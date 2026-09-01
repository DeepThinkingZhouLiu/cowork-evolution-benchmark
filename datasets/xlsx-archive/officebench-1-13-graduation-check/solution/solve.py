#!/usr/bin/env python3
"""OfficeBench 1-13/2 的逐题参考解法。"""

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

from openpyxl import load_workbook

INPUT_NAME = "Thomas_Robinson_Course_List.xlsx"
OUTPUT_NAME = "answer.txt"
EXPECTED_INPUT_SHA256 = "5a495d6f9d3c319207dc1b9ed002672a42eee30c3f25741b31ca4e7df399faeb"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    source = args.input_dir / INPUT_NAME
    if sha256(source) != EXPECTED_INPUT_SHA256:
        raise ValueError("课程记录输入哈希不匹配")
    workbook = load_workbook(source, data_only=True)
    sheet = workbook["Sheet1"]
    cs_count = 0
    math_count = 0
    for row in range(2, sheet.max_row + 1):
        course_id = str(sheet.cell(row, 3).value or "")
        if course_id.startswith("CS "):
            cs_count += 1
        if course_id.startswith("MATH "):
            math_count += 1
    answer = "Yes" if cs_count >= 8 and math_count >= 5 else "No"
    args.output_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / OUTPUT_NAME).write_text(answer + "\n", encoding="utf-8")
    workbook.close()


if __name__ == "__main__":
    main()
