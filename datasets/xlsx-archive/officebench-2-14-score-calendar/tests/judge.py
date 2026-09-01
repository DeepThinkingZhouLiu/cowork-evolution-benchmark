#!/usr/bin/env python3
"""从输入工作簿独立重算前三名并检查离线 iCalendar 交付物。"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

from openpyxl import load_workbook

INPUT_NAME = "score_input.xlsx"
EXPECTED_SHA256 = "a5c8026dc09f38411c6b8783dad74f7f66f08387eeb21b9d8eb7f28a90dd5b6e"
EXPECTED_DATE_TIME = "20200501T100000"
EXPECTED_NAMES = ["Tina", "Rachel", "Sam"]
EXPECTED_FILES = [f"{name}.ics" for name in EXPECTED_NAMES]
TASK_ID = "cowork-evo/officebench-2-14-score-calendar"
RUBRIC_VERSION = "officebench-2-14-score-calendar-v0.1"
JUDGE_VERSION = "officebench-2-14-score-calendar-deterministic-v0.1"
WEIGHTS = {"R001": 0.2, "R002": 0.25, "R003": 0.25, "R004": 0.15, "R005": 0.05, "R006": 0.1}


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def criterion(cid: str, passed: bool, evidence: str) -> dict:
    return {"criterion_id": cid, "score": 1.0 if passed else 0.0, "evidence": evidence, "method": "deterministic"}


def unescape(value: str) -> str:
    return value.replace("\\n", "\n").replace("\\,", ",").replace("\\;", ";").replace("\\\\", "\\")


def parse_event(path: Path) -> dict[str, str]:
    text = path.read_bytes().decode("utf-8")
    if not text.endswith("\r\n") or "BEGIN:VCALENDAR\r\n" not in text or "END:VCALENDAR\r\n" not in text:
        raise ValueError("缺少完整 VCALENDAR 边界或 CRLF")
    if text.count("BEGIN:VEVENT\r\n") != 1 or text.count("END:VEVENT\r\n") != 1:
        raise ValueError("VEVENT 数量不是 1")
    body = text.split("BEGIN:VEVENT\r\n", 1)[1].split("\r\nEND:VEVENT\r\n", 1)[0]
    props: dict[str, str] = {}
    for line in body.split("\r\n"):
        if ":" not in line:
            continue
        key, value = line.split(":", 1)
        props[key.split(";", 1)[0].upper()] = unescape(value)
    required = {"UID", "DTSTAMP", "DTSTART", "SUMMARY", "DESCRIPTION"}
    missing = required - set(props)
    if missing:
        raise ValueError(f"缺少属性：{sorted(missing)}")
    if not re.fullmatch(r"[0-9]{8}T[0-9]{6}Z", props["DTSTAMP"]):
        raise ValueError("DTSTAMP 不是 UTC 时间")
    return props


def top_three(input_path: Path) -> list[tuple[str, float]]:
    workbook = load_workbook(input_path, data_only=True, read_only=True)
    if workbook.sheetnames != ["Sheet1"]:
        raise ValueError(f"工作表结构异常：{workbook.sheetnames}")
    sheet = workbook["Sheet1"]
    headers = [sheet.cell(1, col).value for col in range(1, 4)]
    if headers != ["Name", "midterm1", "midterm2"]:
        raise ValueError(f"表头异常：{headers!r}")
    rows = []
    for row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if row[0] is not None:
            rows.append((row_index, str(row[0]), float(row[1])))
    workbook.close()
    return [(name, score) for _, name, score in sorted(rows, key=lambda item: (-item[2], item[0]))[:3]]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    source_path = Path(__file__).parents[1] / "data/input_files" / INPUT_NAME
    results: list[dict] = []
    try:
        expected = top_three(source_path)
        expected_names = [name for name, _ in expected]
        actual_files = sorted(path.name for path in args.artifact_dir.iterdir() if path.is_file()) if args.artifact_dir.is_dir() else []
        missing = [name for name in EXPECTED_FILES if name not in actual_files]
        parse_errors: list[str] = []
        parsed: dict[str, dict[str, str]] = {}
        for filename in EXPECTED_FILES:
            path = args.artifact_dir / filename
            if not path.is_file():
                continue
            try:
                parsed[filename[:-4]] = parse_event(path)
            except Exception as exc:
                parse_errors.append(f"{filename}: {exc}")

        r001 = not missing and not parse_errors and len(parsed) == 3
        results.append(criterion("R001", r001, f"缺失={missing}; 解析错误={parse_errors}"))
        names_ok = expected_names == EXPECTED_NAMES and sorted(parsed) == sorted(EXPECTED_NAMES) and set(actual_files) == set(EXPECTED_FILES)
        results.append(criterion("R002", names_ok, f"输入前三名={expected_names}; 输出={sorted(parsed)}; 额外文件={sorted(set(actual_files) - set(EXPECTED_FILES))}"))

        timing_ok = all(props.get("DTSTART") == EXPECTED_DATE_TIME for props in parsed.values()) and len(parsed) == 3
        identity_ok = all(name in props.get("SUMMARY", "") and name in props.get("DESCRIPTION", "") for name, props in parsed.items())
        results.append(criterion("R003", timing_ok and identity_ok, f"时间={timing_ok}; 文件与姓名对应={identity_ok}"))

        content_ok = all("Reward meeting" in props.get("SUMMARY", "") and "top-three" in props.get("DESCRIPTION", "") and "midterm1" in props.get("DESCRIPTION", "") for props in parsed.values()) and len(parsed) == 3
        results.append(criterion("R004", content_ok, f"标题/描述语义={content_ok}"))
        native_ok = all(props.get("UID") and props.get("DTSTAMP") for props in parsed.values()) and not parse_errors and len(parsed) == 3
        results.append(criterion("R005", native_ok and set(actual_files) == set(EXPECTED_FILES), f"原生 VEVENT={native_ok}"))
        hash_ok = digest(source_path) == EXPECTED_SHA256 and expected == [("Tina", 100.0), ("Rachel", 100.0), ("Sam", 99.0)]
        results.append(criterion("R006", hash_ok, f"输入 SHA256={digest(source_path)}; 独立前三名={expected}"))
        reward = sum(WEIGHTS[item["criterion_id"]] * item["score"] for item in results)
    except Exception as exc:
        results = [criterion("R001", False, f"Judge 读取失败：{exc}")]
        reward = 0.0
    by_id = {item["criterion_id"]: item["score"] for item in results}
    passed = reward >= 0.7 and all(by_id.get(cid) == 1.0 for cid in ("R001", "R002", "R003"))
    payload = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [digest(path) for path in sorted(args.artifact_dir.glob("*.ics"))] if args.artifact_dir.is_dir() else [], "criterion_results": results, "reward": round(reward, 6), "passed": passed}
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.write_text(f"{reward:.6f}\n", encoding="utf-8")
    raise SystemExit(0 if passed else 1)


if __name__ == "__main__":
    main()
