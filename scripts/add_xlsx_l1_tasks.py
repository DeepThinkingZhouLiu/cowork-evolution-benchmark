from __future__ import annotations

import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from scripts.build_xlsx_batch import OFFICE, TASKS, make_xlsx, office_meta, sha, write

REV = "b978b808667c32b52ce19a67ce1def1de9ae02b7"
OFFICE_URI = f"https://github.com/zlwang-cs/OfficeBench/tree/{REV}"
LICENSE = Path("/tmp/OfficeBench/LICENSE")

def make_golden(src: Path, dst: Path, operation: str) -> None:
    wb = load_workbook(src)
    ws = wb.active
    if operation == "attendance":
        ws.cell(1, 2).value = "Attendance"
        for row in range(2, ws.max_row + 1):
            ws.cell(row, 2).value = "Y"
    elif operation == "highest_salary_duplicate":
        amount, name = max((ws.cell(r, 2).value, ws.cell(r, 1).value) for r in range(2, ws.max_row + 1))
        ws.append([name, amount])
    elif operation == "lowest_budget_add_highest":
        amounts = [ws.cell(r, 2).value for r in range(2, ws.max_row + 1)]
        ws.cell(amounts.index(max(amounts)) + 2, 2).value = max(amounts) + min(amounts)
    elif operation in {"total_salary", "total_budget", "total_shopping_weight"}:
        ws.append(["Total", sum(ws.cell(r, 2).value for r in range(2, ws.max_row + 1))])
    else:
        raise ValueError(operation)
    wb.save(dst)
    wb.close()

def solver_code(input_name: str, output_name: str, operation: str) -> str:
    return f'''from pathlib import Path
import os
from openpyxl import load_workbook
INPUT = {input_name!r}
OUTPUT = {output_name!r}
OPERATION = {operation!r}
def main():
    root = Path(__file__).parents[1]
    src = root / "data/input_files" / INPUT
    out_dir = Path(os.environ.get("OUTPUT_DIR", "/app/output"))
    out_dir.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(src)
    ws = wb.active
    if OPERATION == "attendance":
        ws.cell(1, 2).value = "Attendance"
        for row in range(2, ws.max_row + 1):
            ws.cell(row, 2).value = "Y"
    elif OPERATION == "highest_salary_duplicate":
        amount, name = max((ws.cell(r, 2).value, ws.cell(r, 1).value) for r in range(2, ws.max_row + 1))
        ws.append([name, amount])
    elif OPERATION == "lowest_budget_add_highest":
        amounts = [ws.cell(r, 2).value for r in range(2, ws.max_row + 1)]
        ws.cell(amounts.index(max(amounts)) + 2, 2).value = max(amounts) + min(amounts)
    elif OPERATION in {{"total_salary", "total_budget", "total_shopping_weight"}}:
        ws.append(["Total", sum(ws.cell(r, 2).value for r in range(2, ws.max_row + 1))])
    else:
        raise ValueError(OPERATION)
    wb.save(out_dir / OUTPUT)
    wb.close()
if __name__ == "__main__":
    main()
'''

def judge_code(task_id: str, input_name: str, output_name: str, operation: str) -> str:
    return f'''#!/usr/bin/env python3
from __future__ import annotations
import argparse, hashlib, json, zipfile
from pathlib import Path
from openpyxl import load_workbook
TASK_ID = {task_id!r}
INPUT = {input_name!r}
OUTPUT = {output_name!r}
OPERATION = {operation!r}
WEIGHTS = {{"R001": .10, "R002": .45, "R003": .15, "R004": .20, "R005": .05, "R006": .05, "R007": -.20}}
def criterion(cid, ok, evidence):
    return {{"criterion_id": cid, "score": 1.0 if ok else 0.0, "evidence": evidence, "method": "deterministic"}}
def expected_from_input(wb):
    ws = wb.active
    expected = {{(s.title, cell.coordinate): cell.value for s in wb.worksheets for row in s.iter_rows() for cell in row}}
    if OPERATION == "attendance":
        expected[(ws.title, "B1")] = "Attendance"
        for row in range(2, ws.max_row + 1): expected[(ws.title, f"B{{row}}")] = "Y"
        shape = (ws.max_row, 2)
    elif OPERATION == "highest_salary_duplicate":
        amount, name = max((ws.cell(r, 2).value, ws.cell(r, 1).value) for r in range(2, ws.max_row + 1))
        row = ws.max_row + 1
        expected[(ws.title, f"A{{row}}")], expected[(ws.title, f"B{{row}}")] = name, amount
        shape = (row, 2)
    elif OPERATION == "lowest_budget_add_highest":
        amounts = [ws.cell(r, 2).value for r in range(2, ws.max_row + 1)]
        target = amounts.index(max(amounts)) + 2
        expected[(ws.title, f"B{{target}}")] = max(amounts) + min(amounts)
        shape = (ws.max_row, ws.max_column)
    elif OPERATION in {{"total_salary", "total_budget", "total_shopping_weight"}}:
        row = ws.max_row + 1
        expected[(ws.title, f"A{{row}}")] = "Total"
        expected[(ws.title, f"B{{row}}")] = sum(ws.cell(r, 2).value for r in range(2, ws.max_row + 1))
        shape = (row, 2)
    else: raise ValueError(OPERATION)
    return expected, shape
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--artifact-dir", type=Path, required=True)
    ap.add_argument("--result", type=Path, required=True)
    ap.add_argument("--reward-file", type=Path, required=True)
    a = ap.parse_args()
    root = Path(__file__).parents[1]; src = root / "data/input_files" / INPUT; out = a.artifact_dir / OUTPUT
    results = []; native = out.is_file() and zipfile.is_zipfile(out)
    results.append(criterion("R001", native, "输出存在且为可解析 OOXML 工作簿"))
    target_ok = preservation_ok = formula_ok = parse_ok = False; mismatch_count = 0
    if native:
        try:
            wi = load_workbook(src, data_only=False); wo = load_workbook(out, data_only=False)
            expected, shape = expected_from_input(wi); parse_ok = True
            target_ok = wi.sheetnames == wo.sheetnames and (wo.active.max_row, wo.active.max_column) == shape
            if target_ok:
                for (sheet, coord), value in expected.items():
                    if wo[sheet][coord].value != value: mismatch_count += 1
                target_ok = mismatch_count == 0; preservation_ok = target_ok; formula_ok = True
            wi.close(); wo.close()
            results.append(criterion("R002", target_ok, f"从输入独立重算，错误单元格={{mismatch_count}}"))
            results.append(criterion("R003", formula_ok, "目标为可编辑单元格值，未要求隐藏计算"))
            results.append(criterion("R004", preservation_ok, "逐格检查非目标内容和原工作表结构"))
            results.append(criterion("R005", parse_ok, "openpyxl 可解析工作簿"))
            input_hash_ok = hashlib.sha256(src.read_bytes()).hexdigest() == hashlib.sha256((root / "environment/input_files" / INPUT).read_bytes()).hexdigest()
            results.append(criterion("R006", input_hash_ok, "输入副本 SHA256 一致"))
            results.append(criterion("R007", False, "未检测到空白或硬编码替代品"))
        except Exception as exc:
            results.extend([criterion("R002", False, f"读取或重算失败: {{exc}}"), criterion("R003", False, "读取失败"), criterion("R004", False, "读取失败"), criterion("R005", False, "读取失败"), criterion("R006", False, "读取失败"), criterion("R007", True, "损坏输出触发惩罚")])
    else:
        results.extend([criterion("R002", False, "缺少有效输出"), criterion("R003", False, "缺少有效输出"), criterion("R004", False, "缺少有效输出"), criterion("R005", False, "缺少有效输出"), criterion("R006", False, "缺少有效输出"), criterion("R007", True, "缺少有效输出触发惩罚")])
    vals = {{x["criterion_id"]: x["score"] for x in results}}; reward = sum(WEIGHTS[k] * vals.get(k, 0) for k in WEIGHTS)
    passed = vals.get("R001", 0) == 1 and vals.get("R002", 0) == 1 and reward >= .7
    payload = {{"task_id": TASK_ID, "rubric_version": "cowork-evo-rubric-v0.1", "judge_version": "officebench-l1-independent-v0.1", "criterion_results": results, "reward": reward, "passed": passed}}
    a.result.parent.mkdir(parents=True, exist_ok=True); a.result.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    a.reward_file.parent.mkdir(parents=True, exist_ok=True); a.reward_file.write_text(f"{{reward:.6f}}\\n")
    raise SystemExit(0 if passed else 1)
if __name__ == "__main__":
    main()
'''

TASKS_TO_ADD = [
    ("officebench-1-6-attendance", "1-6", 1, "class_member.xlsx", "class_attendance.xlsx", "attendance", "在 class_member.xlsx 中新增 Attendance 列，为每位学生填入 Y。", "class attendance", "教育管理", "课程助理"),
    ("officebench-1-10-highest-salary-duplicate", "1-10", 2, "salary.xlsx", "salary_highest_duplicate.xlsx", "highest_salary_duplicate", "读取 salary.xlsx，找出 amount 最高的记录，并将该整行复制到表格底部。", "salary row duplication", "薪酬管理", "薪酬助理"),
    ("officebench-1-10-lowest-budget-add-highest", "1-10", 3, "company_budget.xlsx", "budget_lowest_to_highest.xlsx", "lowest_budget_add_highest", "读取 company_budget.xlsx，将最低 amount 加到最高 amount 所在行，保留其他内容。", "budget arithmetic", "财务预算", "预算助理"),
    ("officebench-1-9-total-salary", "1-9", 2, "salary.xlsx", "total_salary.xlsx", "total_salary", "读取 salary.xlsx，在底部追加 Total 行，填入所有 amount 的总和。", "salary aggregation", "薪酬管理", "薪酬助理"),
    ("officebench-1-9-total-budget", "1-9", 3, "company_budget.xlsx", "total_budget.xlsx", "total_budget", "读取 company_budget.xlsx，在底部追加 Total 行，填入所有 amount 的总和。", "budget aggregation", "财务预算", "预算助理"),
    ("officebench-1-9-total-shopping-weight", "1-9", 4, "shopping_list.xlsx", "total_shopping_weight.xlsx", "total_shopping_weight", "读取 shopping_list.xlsx，在底部追加 Total 行，填入所有商品 amount 的总和。", "shopping aggregation", "采购管理", "采购助理"),
]

def build_one(item) -> None:
    task_id, group, subtask, input_file, output_file, operation, instruction, family, sector, occupation = item
    src = OFFICE / f"tasks/{group}/testbed/data/{input_file}"
    golden_tmp = Path("/tmp") / f"{task_id}-golden.xlsx"
    make_golden(src, golden_tmp, operation)
    meta = office_meta(task_id, "basic", f"改造自 OfficeBench {group}/{subtask}：{instruction}", f"请打开输入文件 {input_file}。{instruction} 将结果保存为 /app/output/{output_file}。保留未涉及的单元格、工作表结构和格式，交付可继续编辑的原生 Excel 工作簿，不要联网。", family, output_file, "OfficeBench", f"{OFFICE_URI}/tasks/{group}", REV, "L1", sector, occupation, f"officebench-{task_id}", 10)
    meta["operation_types"] = ["edit"]; meta["capability_ids"] = ["K3", "K7", "K12"]; meta["supplier_capability_ids"] = ["C7", "C8", "C12"]
    groups = {"attendance": "officebench-class-attendance", "highest_salary_duplicate": "officebench-salary-basic", "total_salary": "officebench-salary-basic", "lowest_budget_add_highest": "officebench-budget-basic", "total_budget": "officebench-budget-basic", "total_shopping_weight": "officebench-shopping-basic"}
    meta["leakage_group_id"] = groups[operation]
    make_xlsx(meta, src, golden_tmp, input_file, output_file)
    root = TASKS / task_id
    write(root / "solution/solve.py", solver_code(input_file, output_file, operation))
    write(root / "tests/judge.py", judge_code(f"cowork-evo/{task_id}", input_file, output_file, operation))
    shutil.copy2(OFFICE / f"tasks/{group}/subtasks/{subtask}.json", root / "adaptation" / "upstream-subtask.json")
    shutil.copy2(LICENSE, root / "adaptation" / "LICENSE-OfficeBench")
    digest = sha(root / "data/input_files" / input_file)
    write(root / "adaptation/README.md", f'''# {task_id}

状态：pending_human_review；difficulty_tier=basic 是 L1 作者假设，calibration_status=uncalibrated，review_status=pending_human_review。

## 来源与证据

- 来源：OfficeBench tasks/{group}/subtasks/{subtask}.json，固定 revision {REV}；原始子任务 JSON 与 Apache-2.0 许可证保留在本目录。
- 输入：data/input_files/{input_file}，SHA256 {digest}，来自上游 testbed/data/{input_file}。
- 独立真值：Judge 从 solver-visible 输入重算目标，不读取 Golden；预期结果为输入驱动的单工作簿编辑。

## 改造边界

仅将上游任务改造为离线本地输入/输出，明确原生可编辑工作簿和非目标保持要求；未增加输入之外的业务事实。参考解法同样从输入计算并写出结果。

## 验证记录

作者侧应执行 validate-task、validate-task-assets、Golden/fresh replay，以及空输出、损坏输出、未修改输入和目标错误负例。Judge 覆盖文件完整性、输入驱动目标、非目标保持、OOXML 可解析性和输入哈希。尚未执行无辅助模型 rollout，最终难度和发布状态待人工复核。
''')

if __name__ == "__main__":
    for item in TASKS_TO_ADD:
        build_one(item)
