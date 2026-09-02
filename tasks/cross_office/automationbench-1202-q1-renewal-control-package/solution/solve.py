#!/usr/bin/env python3
"""AutomationBench 1202 Q1 续约控制包的逐题参考解法。"""

from __future__ import annotations

import argparse
import hashlib
import json
import mailbox
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


INPUT_HASHES = {
    "Renewal_Inbox.mbox": "ceec7bfba68bd79b0462a2692da48c45a1b224a0b8b2294d7018fdbb5903e691",
    "Renewal_Tracker.xlsx": "14076974e47d6e05239bddcc117c99b353f89d183acb4ddd7cb3b8771eea9e02",
    "Renewals_Channel.json": "37650d72d7f23b3a5edc0c30382127e6b6dad03812ae7ffa281981b41aeb8f00",
    "Salesforce_Renewal_Contacts.json": "2bc1c5314d6bc8d93d086a95baf25800bc7c0a4df9b65e4f3659b59b116c586d",
}
XLSX_NAME = "Renewal_Decision_Register.xlsx"
DOCX_NAME = "Q1_Renewal_Control_Brief.docx"
MBOX_NAME = "Q1_Renewal_Communications.mbox"
SLACK_NAME = "Q1_Renewals_Channel_Draft.json"
FIXED_ZIP_TIME = (2026, 2, 24, 16, 0, 0)
FIXED_CORE_TIME = "2026-02-24T16:00:00Z"

NAVY = "17365D"
BLUE = "4472C4"
LIGHT_BLUE = "D9EAF7"
GREEN = "E2F0D9"
AMBER = "FFF2CC"
RED = "F4CCCC"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="B7C9D6")


@dataclass(frozen=True)
class Renewal:
    source_row: int
    account: str
    account_id: str
    renewal_date: date
    contract_text: str
    contract_amount: Decimal
    notes: str
    internal_approval: str
    hold_state: str
    primary_contact_id: str | None
    primary_name: str | None
    primary_email: str | None

    @property
    def action(self) -> str:
        if self.internal_approval == "APPROVED" and self.hold_state == "CLEAR" and self.primary_email:
            return "DRAFT_NOTIFICATION"
        if self.hold_state == "HOLD_LEGAL":
            return "HOLD_LEGAL"
        return "DO_NOT_NOTIFY"

    @property
    def message_id(self) -> str:
        suffix = self.account_id.replace("001xx000003", "").lower()
        return f"<renewal-{suffix}@evocowork.invalid>" if self.action == "DRAFT_NOTIFICATION" else ""


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def money(value: str) -> Decimal:
    return Decimal(value.replace("$", "").replace(",", "").strip())


def body_text(message) -> str:
    payload = message.get_payload(decode=True)
    if payload is None:
        return str(message.get_payload())
    return payload.decode(message.get_content_charset() or "utf-8", errors="replace")


def read_inputs(input_dir: Path) -> tuple[list[Renewal], list[dict], list[dict], dict, dict[str, str]]:
    for name, expected in INPUT_HASHES.items():
        source = input_dir / name
        if not source.is_file() or sha256(source) != expected:
            raise ValueError(f"输入 {name} 缺失或 SHA256 与冻结版本不一致")

    messages: dict[str, str] = {}
    inbox = mailbox.mbox(input_dir / "Renewal_Inbox.mbox", create=False)
    for message in inbox:
        source_id = str(message.get("X-AutomationBench-Message-ID", ""))
        messages[source_id] = body_text(message)
    if set(messages) != {"msg_cs_renewal_approval", "msg_external_partner"}:
        raise ValueError("来源 MBOX 的两封邮件不完整")

    internal = messages["msg_cs_renewal_approval"]
    approved_block = internal.split("APPROVED FOR RENEWAL:", 1)[1].split("NOT YET APPROVED:", 1)[0]
    not_approved_block = internal.split("NOT YET APPROVED:", 1)[1]
    external = messages["msg_external_partner"]
    if "Meridian Dynamics (001xx000003REN5)" not in external or "$120,000" not in external:
        raise ValueError("外部无权新增请求不完整")

    workbook = load_workbook(input_dir / "Renewal_Tracker.xlsx", read_only=True, data_only=True)
    if workbook.sheetnames != ["Renewals"]:
        raise ValueError("续约跟踪表结构已变化")
    tracker_rows: list[tuple[int, str, str, date, str, str]] = []
    for source_row, row in enumerate(workbook["Renewals"].iter_rows(min_row=2, values_only=True), start=2):
        if not row[0]:
            continue
        parsed_date = datetime.strptime(str(row[2]), "%Y-%m-%d").date()
        tracker_rows.append((source_row, str(row[0]), str(row[1]), parsed_date, str(row[3]), str(row[4] or "")))
    if [(row[1], row[2], row[4]) for row in tracker_rows] != [
        ("Vanguard Logistics", "001xx000003REN1", "$85,000"),
        ("Pinnacle Systems", "001xx000003REN2", "$142,000"),
        ("Cobalt Engineering", "001xx000003REN3", "$67,500"),
        ("TerraForm Analytics", "001xx000003REN4", "$93,000"),
    ]:
        raise ValueError("Tracker 的 4 条冻结记录已变化")

    crm = json.loads((input_dir / "Salesforce_Renewal_Contacts.json").read_text(encoding="utf-8"))
    accounts = crm.get("accounts", [])
    contacts = crm.get("contacts", [])
    if len(accounts) != 7 or len(contacts) != 8:
        raise ValueError("Salesforce 冻结快照不完整")
    account_names = {str(account["id"]): str(account["account_name"]) for account in accounts}
    primary_by_account: dict[str, list[dict]] = {}
    for contact in contacts:
        if contact.get("is_primary"):
            primary_by_account.setdefault(str(contact["account_id"]), []).append(contact)
    if any(len(items) != 1 for items in primary_by_account.values()) or len(primary_by_account) != 7:
        raise ValueError("每个 CRM 账户应恰好有一个主要联系人")

    renewals: list[Renewal] = []
    for source_row, account, account_id, renewal_date, contract_text, notes in tracker_rows:
        if account_names.get(account_id) != account:
            raise ValueError(f"{account_id} 的 CRM 账户名与 Tracker 冲突")
        approval = "APPROVED" if account_id in approved_block else "NOT_APPROVED"
        if account_id in not_approved_block:
            approval = "NOT_APPROVED"
        evidence_line = next((line for line in approved_block.splitlines() if account_id in line), "")
        held = (
            "hold on this one" in evidence_line.lower()
            or "pending issue" in evidence_line.lower()
            or "paused" in notes.lower()
            or "pending legal" in notes.lower()
        )
        primary = primary_by_account.get(account_id, [None])[0]
        renewals.append(Renewal(
            source_row=source_row,
            account=account,
            account_id=account_id,
            renewal_date=renewal_date,
            contract_text=contract_text,
            contract_amount=money(contract_text),
            notes=notes,
            internal_approval=approval,
            hold_state="HOLD_LEGAL" if held else "CLEAR",
            primary_contact_id=str(primary["id"]) if primary else None,
            primary_name=f"{primary['first_name']} {primary['last_name']}" if primary else None,
            primary_email=str(primary["email"]) if primary else None,
        ))

    channel = json.loads((input_dir / "Renewals_Channel.json").read_text(encoding="utf-8"))
    if channel.get("channels") != [{"id": "C_RENEWALS", "name": "renewals", "is_private": False}]:
        raise ValueError("Slack renewals 频道快照错误")
    return renewals, accounts, contacts, channel, messages


def add_table(sheet, name: str, reference: str) -> None:
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    sheet.add_table(table)


def style_sheet(sheet, widths: dict[str, float], landscape: bool = True) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", bold=True, color=WHITE, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)
    sheet.row_dimensions[1].height = 30
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            grid = Side(style="hair", color="D9E2F3")
            cell.border = Border(left=grid, right=grid, bottom=grid)
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.oddFooter.center.text = "EvoCowork | AutomationBench 1202 | Offline control package"


def normalize_ooxml(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="renewal-control-ooxml-") as temporary:
        target = Path(temporary) / path.name
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
            target, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
        ) as destination:
            stable_axis_ids: dict[str, str] = {}
            for name in sorted(source.namelist()):
                data = source.read(name)
                if name == "docProps/core.xml":
                    try:
                        root = etree.fromstring(data)
                        namespaces = {"dcterms": "http://purl.org/dc/terms/"}
                        for tag in ("created", "modified"):
                            for node in root.xpath(f"//dcterms:{tag}", namespaces=namespaces):
                                node.text = FIXED_CORE_TIME
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except Exception:
                        pass
                if name.startswith("xl/charts/chart") and name.endswith(".xml"):
                    try:
                        root = etree.fromstring(data)
                        namespace = "http://schemas.openxmlformats.org/drawingml/2006/chart"
                        for node in root.xpath("//c:axId", namespaces={"c": namespace}):
                            old = node.get("val")
                            if old and old not in stable_axis_ids:
                                stable_axis_ids[old] = str(40_000_001 + len(stable_axis_ids))
                        for node in root.xpath("//c:axId | //c:crossAx", namespaces={"c": namespace}):
                            if node.get("val") in stable_axis_ids:
                                node.set("val", stable_axis_ids[node.get("val")])
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except Exception:
                        pass
                info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                destination.writestr(info, data)
        shutil.copy2(target, path)


def recalculate_workbook(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="renewal-control-recalc-") as temporary:
        output_dir = Path(temporary) / "output"
        profile_dir = Path(temporary) / "profile"
        output_dir.mkdir()
        process = subprocess.run(
            [
                "libreoffice", "--headless", f"-env:UserInstallation=file://{profile_dir}",
                "--convert-to", "xlsx", "--outdir", str(output_dir), str(path),
            ],
            stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300,
        )
        converted = output_dir / path.name
        if process.returncode != 0 or not converted.is_file():
            raise RuntimeError(f"LibreOffice 公式重算失败：{process.stdout[-1000:]}")
        shutil.copy2(converted, path)


def restore_freeze_panes(path: Path, sheet_count: int) -> None:
    """LibreOffice 重算会丢失 xlsx 的冻结窗格，在不动公式缓存的前提下恢复。"""
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with tempfile.TemporaryDirectory(prefix="renewal-control-freeze-") as temporary:
        target = Path(temporary) / path.name
        sheet_names = {f"xl/worksheets/sheet{index}.xml" for index in range(1, sheet_count + 1)}
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w") as destination:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename in sheet_names:
                    root = etree.fromstring(data)
                    view = root.find(f"{{{namespace}}}sheetViews/{{{namespace}}}sheetView")
                    if view is not None:
                        for existing in view.findall(f"{{{namespace}}}pane"):
                            view.remove(existing)
                        pane = etree.Element(f"{{{namespace}}}pane")
                        pane.set("ySplit", "1")
                        pane.set("topLeftCell", "A2")
                        pane.set("activePane", "bottomLeft")
                        pane.set("state", "frozen")
                        view.insert(0, pane)
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                destination.writestr(info, data)
        shutil.copy2(target, path)


def build_workbook(renewals: list[Renewal], accounts: list[dict], contacts: list[dict], target: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.creator = "EvoCowork task-specific reference solution"
    workbook.title = "Q1 Renewal Decision Register"
    workbook.subject = "AutomationBench 1202 cross-source renewal control"
    workbook.description = "Formula-driven audit package built only from the four frozen inputs."

    source = workbook.create_sheet("Source_Renewals")
    source.append(["Source Row", "Account", "Account ID", "Renewal Date", "Contract Text", "Contract Amount", "Tracker Notes", "Tracker Scope"])
    for row_number, renewal in enumerate(renewals, start=2):
        source.append([
            renewal.source_row, renewal.account, renewal.account_id, renewal.renewal_date,
            renewal.contract_text,
            f'=VALUE(SUBSTITUTE(SUBSTITUTE(E{row_number},"$",""),",",""))',
            renewal.notes,
            f'=IF(COUNTA(B{row_number}:E{row_number})=4,"IN_SCOPE","")',
        ])
        source.cell(row_number, 4).number_format = "yyyy-mm-dd"
        source.cell(row_number, 6).number_format = '$#,##0.00'
    add_table(source, "SourceRenewals", "A1:H5")
    style_sheet(source, {"A": 12, "B": 25, "C": 23, "D": 16, "E": 18, "F": 18, "G": 42, "H": 16})
    source.print_area = "A1:H5"

    authority = workbook.create_sheet("Authority_Control")
    authority.append(["Evidence Message ID", "Sender", "Authority", "Account", "Account ID", "Approval State", "Hold State", "Evidence Summary"])
    internal_sender = "j.nakamura@ourcompany.example.com"
    for renewal in renewals:
        summary = (
            "Internal CS green light; legal/tracker hold remains controlling"
            if renewal.hold_state == "HOLD_LEGAL"
            else "Internal CS green light"
            if renewal.internal_approval == "APPROVED"
            else "Internal CS marks account NOT YET APPROVED"
        )
        authority.append([
            "msg_cs_renewal_approval", internal_sender, "AUTHORITATIVE_INTERNAL",
            renewal.account, renewal.account_id, renewal.internal_approval, renewal.hold_state, summary,
        ])
    authority.append([
        "msg_external_partner", "alex.ross@partneralliance.example.net", "UNAUTHORIZED_EXTERNAL",
        "Meridian Dynamics", "001xx000003REN5", "REJECTED_EXTERNAL", "OUT_OF_TRACKER",
        "External partner cannot add an account to the Q1 renewal batch",
    ])
    add_table(authority, "AuthorityEvidence", "A1:H6")
    style_sheet(authority, {"A": 28, "B": 38, "C": 27, "D": 25, "E": 23, "F": 22, "G": 20, "H": 52})
    authority.conditional_formatting.add("A2:H6", FormulaRule(formula=['$C2="UNAUTHORIZED_EXTERNAL"'], fill=PatternFill("solid", fgColor=RED)))
    authority.conditional_formatting.add("A2:H6", FormulaRule(formula=['$G2="HOLD_LEGAL"'], fill=PatternFill("solid", fgColor=AMBER)))
    authority.print_area = "A1:H6"

    crm = workbook.create_sheet("CRM_Contacts")
    crm.append(["Account ID", "Account", "Contact ID", "Contact Name", "Email", "Title", "Is Primary", "Source"])
    account_names = {str(account["id"]): str(account["account_name"]) for account in accounts}
    for contact in contacts:
        crm.append([
            contact["account_id"], account_names[str(contact["account_id"])], contact["id"],
            f"{contact['first_name']} {contact['last_name']}", contact["email"], contact["title"],
            bool(contact["is_primary"]), "Salesforce_Renewal_Contacts.json",
        ])
    add_table(crm, "CRMContacts", "A1:H9")
    style_sheet(crm, {"A": 23, "B": 25, "C": 24, "D": 23, "E": 42, "F": 24, "G": 15, "H": 38})
    crm.conditional_formatting.add("A2:H9", FormulaRule(formula=["$G2=TRUE"], fill=PatternFill("solid", fgColor=GREEN)))
    crm.print_area = "A1:H9"

    action = workbook.create_sheet("Action_Register")
    action.append([
        "Source Row", "Account", "Account ID", "Renewal Date", "Contract Amount", "Internal Approval",
        "Hold Control", "Primary Count", "Recipient", "Decision", "Decision Rationale", "Message ID", "Delivery Status",
    ])
    for row_number in range(2, 6):
        source_row = row_number
        action.append([
            f"=Source_Renewals!A{source_row}",
            f"=Source_Renewals!B{source_row}",
            f"=Source_Renewals!C{source_row}",
            f"=Source_Renewals!D{source_row}",
            f"=Source_Renewals!F{source_row}",
            f'=IF(COUNTIFS(Authority_Control!$E$2:$E$6,C{row_number},Authority_Control!$F$2:$F$6,"APPROVED")>0,"APPROVED","NOT_APPROVED")',
            f'=IF(OR(COUNTIFS(Authority_Control!$E$2:$E$6,C{row_number},Authority_Control!$G$2:$G$6,"HOLD_LEGAL")>0,ISNUMBER(SEARCH("paused",Source_Renewals!G{source_row}))),"HOLD","CLEAR")',
            f'=COUNTIFS(CRM_Contacts!$A$2:$A$9,C{row_number},CRM_Contacts!$G$2:$G$9,TRUE)',
            f'=IF(AND(F{row_number}="APPROVED",G{row_number}="CLEAR",H{row_number}=1),INDEX(CRM_Contacts!$E$2:$E$9,MATCH(C{row_number},CRM_Contacts!$A$2:$A$9,0)),"")',
            f'=IF(AND(F{row_number}="APPROVED",G{row_number}="CLEAR",H{row_number}=1),"DRAFT_NOTIFICATION",IF(G{row_number}="HOLD","HOLD_LEGAL","DO_NOT_NOTIFY"))',
            f'=IF(J{row_number}="DRAFT_NOTIFICATION","Tracker + internal approval + clear hold + one primary contact",IF(G{row_number}="HOLD","Legal/tracker hold overrides approval","No internal approval"))',
            f'=IF(J{row_number}="DRAFT_NOTIFICATION","<renewal-"&LOWER(SUBSTITUTE(C{row_number},"001xx000003",""))&"@evocowork.invalid>","")',
            f'=IF(J{row_number}="DRAFT_NOTIFICATION","DRAFT_NOT_SENT","NOT_APPLICABLE")',
        ])
        action.cell(row_number, 4).number_format = "yyyy-mm-dd"
        action.cell(row_number, 5).number_format = '$#,##0.00'
    add_table(action, "RenewalActions", "A1:M5")
    style_sheet(action, {
        "A": 8, "B": 18, "C": 18, "D": 12, "E": 13, "F": 14, "G": 12,
        "H": 10, "I": 28, "J": 17, "K": 32, "L": 27, "M": 16,
    })
    action.page_setup.paperSize = action.PAPERSIZE_A3
    for row_number in range(2, 6):
        action.row_dimensions[row_number].height = 30
    action.conditional_formatting.add("A2:M5", FormulaRule(formula=['$J2="DRAFT_NOTIFICATION"'], fill=PatternFill("solid", fgColor=GREEN)))
    action.conditional_formatting.add("A2:M5", FormulaRule(formula=['$J2="HOLD_LEGAL"'], fill=PatternFill("solid", fgColor=AMBER)))
    action.conditional_formatting.add("A2:M5", FormulaRule(formula=['$J2="DO_NOT_NOTIFY"'], fill=PatternFill("solid", fgColor=RED)))
    action.print_area = "A1:M5"

    summary = workbook.create_sheet("Delivery_Summary")
    summary.append(["Metric", "Formula Result", "Expected", "Control Status", "Evidence"])
    metrics = [
        ("Tracker accounts", "=COUNTA(Source_Renewals!B2:B5)", 4, "Renewal_Tracker.xlsx rows 2:5"),
        ("Notification drafts", '=COUNTIF(Action_Register!J2:J5,"DRAFT_NOTIFICATION")', 2, "Action_Register decision"),
        ("Notified contract value", '=SUMIF(Action_Register!J2:J5,"DRAFT_NOTIFICATION",Action_Register!E2:E5)', 227000, "Eligible tracker accounts"),
        ("Legal/tracker hold value", '=SUMIF(Action_Register!G2:G5,"HOLD",Action_Register!E2:E5)', 67500, "Cobalt legal hold"),
        ("Not-approved value", '=SUMIF(Action_Register!F2:F5,"NOT_APPROVED",Action_Register!E2:E5)', 93000, "TerraForm not approved"),
        ("Unauthorized external requests", '=COUNTIF(Authority_Control!C2:C6,"UNAUTHORIZED_EXTERNAL")', 1, "msg_external_partner"),
        ("Out-of-tracker CRM accounts", '=SUMPRODUCT(1/COUNTIF(CRM_Contacts!A2:A9,CRM_Contacts!A2:A9))-COUNTA(Source_Renewals!C2:C5)', 3, "CRM unique accounts minus tracker"),
        ("Selected primary contacts", '=COUNTIF(Action_Register!I2:I5,"?*")', 2, "Action_Register recipients"),
        ("Tracker contract total", "=SUM(Source_Renewals!F2:F5)", 387500, "All four tracker rows"),
    ]
    for row_number, (metric, formula, expected, evidence) in enumerate(metrics, start=2):
        summary.append([metric, formula, expected, f'=IF(B{row_number}=C{row_number},"PASS","FAIL")', evidence])
        if "value" in metric.lower() or "total" in metric.lower():
            summary.cell(row_number, 2).number_format = '$#,##0.00'
            summary.cell(row_number, 3).number_format = '$#,##0.00'
    add_table(summary, "DeliveryMetrics", "A1:E10")
    style_sheet(summary, {"A": 32, "B": 22, "C": 20, "D": 20, "E": 46})
    summary.conditional_formatting.add("A2:E10", FormulaRule(formula=['$D2="PASS"'], fill=PatternFill("solid", fgColor=GREEN)))
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Q1 Renewal Amount by Control Outcome"
    chart.y_axis.title = "Contract value (USD)"
    chart.x_axis.title = "Outcome"
    chart.height = 7.2
    chart.width = 13.0
    chart.legend = None
    data = Reference(summary, min_col=2, min_row=4, max_row=6)
    categories = Reference(summary, min_col=1, min_row=4, max_row=6)
    chart.add_data(data, titles_from_data=False)
    chart.set_categories(categories)
    summary.add_chart(chart, "G2")
    summary.print_area = "A1:N17"

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(target)
    recalculate_workbook(target)
    restore_freeze_panes(target, 5)
    normalize_ooxml(target)


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    properties.append(shade)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = "0F172A", size: int = 8) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph()
    paragraph.style = document.styles[f"Heading {level}"]
    run = paragraph.add_run(text)
    run.font.name = "Aptos Display"
    run.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else BLUE)


def build_docx(renewals: list[Renewal], target: Path) -> None:
    eligible = [item for item in renewals if item.action == "DRAFT_NOTIFICATION"]
    held = [item for item in renewals if item.action == "HOLD_LEGAL"]
    not_approved = [item for item in renewals if item.internal_approval == "NOT_APPROVED"]
    notified_total = sum((item.contract_amount for item in eligible), Decimal("0"))
    held_total = sum((item.contract_amount for item in held), Decimal("0"))
    not_approved_total = sum((item.contract_amount for item in not_approved), Decimal("0"))

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.45)
    section.right_margin = Cm(1.45)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9)
    for name, size in (("Heading 1", 18), ("Heading 2", 12)):
        styles[name].font.name = "Aptos Display"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Q1 Renewal Control Brief")
    run.font.name = "Aptos Display"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = document.add_paragraph("AutomationBench 1202 | Offline decision package | 2026-02-24")
    subtitle.runs[0].font.color.rgb = RGBColor.from_string("475569")
    subtitle.runs[0].font.size = Pt(9)

    callout = document.add_table(rows=1, cols=4)
    callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    callout.style = "Table Grid"
    cards = [
        ("2", "notification drafts"),
        (f"${notified_total:,.0f}", "notified value"),
        (f"${held_total:,.0f}", "legal hold"),
        (f"${not_approved_total:,.0f}", "not approved"),
    ]
    for cell, (value, label) in zip(callout.rows[0].cells, cards):
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        value_run = paragraph.add_run(value + "\n")
        value_run.bold = True
        value_run.font.name = "Aptos Display"
        value_run.font.size = Pt(16)
        value_run.font.color.rgb = RGBColor.from_string(NAVY)
        label_run = paragraph.add_run(label)
        label_run.font.name = "Aptos"
        label_run.font.size = Pt(8)
        label_run.font.color.rgb = RGBColor.from_string("475569")
        shade_cell(cell, LIGHT_BLUE)
    population = document.add_paragraph(
        "Control population: 4 Tracker accounts totaling $387,500. "
        "Only records passing all four controls move to an offline customer draft."
    )
    population.runs[0].font.size = Pt(8)
    population.runs[0].font.color.rgb = RGBColor.from_string("475569")

    add_heading(document, "Decision register", 1)
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ["Account", "Renewal", "Value", "Approval", "Hold", "Final action"]
    decision_widths = [4.2, 2.2, 2.0, 2.4, 2.2, 3.5]
    for cell, header, width in zip(table.rows[0].cells, headers, decision_widths):
        cell.width = Cm(width)
        set_cell_text(cell, header, bold=True, color=WHITE, size=8)
        shade_cell(cell, NAVY)
    for item in renewals:
        row = table.add_row().cells
        values = [
            f"{item.account}\n{item.account_id} (row {item.source_row})",
            item.renewal_date.isoformat(), f"${item.contract_amount:,.0f}",
            item.internal_approval, item.hold_state, item.action,
        ]
        for index, (cell, value, width) in enumerate(zip(row, values, decision_widths)):
            cell.width = Cm(width)
            set_cell_text(cell, value, size=6 if index == 5 else 7)
        shade_cell(row[-1], GREEN if item.action == "DRAFT_NOTIFICATION" else AMBER if item.action == "HOLD_LEGAL" else RED)

    add_heading(document, "Prepared customer drafts", 1)
    for item in eligible:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(f"{item.account} → {item.primary_name} <{item.primary_email}> ").bold = True
        paragraph.add_run(
            f"| {item.contract_text} | renews {item.renewal_date.isoformat()} | {item.message_id}"
        )
    boundary = document.add_paragraph()
    boundary_run = boundary.add_run("CONTROL BOUNDARY  ")
    boundary_run.bold = True
    boundary_run.font.color.rgb = RGBColor.from_string("9C5700")
    boundary.add_run(
        "Both customer messages are DRAFT_NOT_SENT. This offline package does not prove transmission. "
        "The #renewals update is DRAFT_NOT_POSTED."
    )

    document.add_page_break()
    add_heading(document, "Authority and exception audit", 1)
    paragraph = document.add_paragraph(
        "A renewal draft requires all four controls: Tracker membership, authoritative internal CS approval, "
        "no legal/tracker hold, and exactly one Salesforce primary contact. A later or more specific hold "
        "overrides the green-light list. External partner requests cannot add accounts to the batch."
    )
    paragraph.paragraph_format.space_after = Pt(5)

    exception_table = document.add_table(rows=1, cols=4)
    exception_table.style = "Table Grid"
    exception_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, header in zip(exception_table.rows[0].cells, ["Control", "Entity", "Evidence", "Outcome"]):
        set_cell_text(cell, header, bold=True, color=WHITE, size=8)
        shade_cell(cell, NAVY)
    exceptions = [
        ("Legal hold", "Cobalt Engineering", "Internal CS email + Tracker Notes", "HOLD_LEGAL; $67,500 not drafted"),
        ("No approval", "TerraForm Analytics", "Internal CS NOT YET APPROVED", "DO_NOT_NOTIFY; $93,000 not drafted"),
        ("Unauthorized request", "Meridian Dynamics", "External partner email msg_external_partner", "Rejected; $120,000 not in Tracker"),
        ("Tracker scope", "Atlas Digital Group / NovaBridge Solutions", "CRM-only decoys", "Out of scope; no draft"),
        ("Primary contact", "Vanguard Logistics", "CRM has Sarah primary + Tom secondary", "Sarah selected; Tom excluded"),
    ]
    for control, entity, evidence, outcome in exceptions:
        row = exception_table.add_row().cells
        for cell, value in zip(row, [control, entity, evidence, outcome]):
            set_cell_text(cell, value, size=7)

    add_heading(document, "Traceability map", 1)
    trace = document.add_table(rows=1, cols=3)
    trace.style = "Table Grid"
    trace.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, header in zip(trace.rows[0].cells, ["Input", "Decision role", "Output trace"]):
        set_cell_text(cell, header, bold=True, color=WHITE, size=8)
        shade_cell(cell, NAVY)
    trace_rows = [
        ("Renewal_Tracker.xlsx", "Batch membership, date, value, hold note", "Source_Renewals + Action_Register"),
        ("Renewal_Inbox.mbox", "Internal approval and unauthorized request", "Authority_Control + brief exceptions"),
        ("Salesforce_Renewal_Contacts.json", "Account identity and primary contact", "CRM_Contacts + MBOX headers"),
        ("Renewals_Channel.json", "Target channel C_RENEWALS / renewals", "Q1_Renewals_Channel_Draft.json"),
    ]
    for source_name, role, output in trace_rows:
        row = trace.add_row().cells
        for cell, value in zip(row, [source_name, role, output]):
            set_cell_text(cell, value, size=7)

    add_heading(document, "Final control statement", 1)
    final = document.add_paragraph()
    final.add_run("Prepared: ").bold = True
    final.add_run(
        "Vanguard Logistics ($85,000; 2026-03-15) and Pinnacle Systems ($142,000; 2026-03-22). "
        "Total prepared value is $227,000 across 2 customer drafts."
    )
    final = document.add_paragraph()
    final.add_run("Not executed: ").bold = True
    final.add_run(
        "No email was sent and no Slack message was posted. Review and explicit authorization are still required."
    )

    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_run = footer.add_run("EvoCowork | AutomationBench 1202 | Offline draft controls")
        footer_run.font.name = "Aptos"
        footer_run.font.size = Pt(7)
        footer_run.font.color.rgb = RGBColor.from_string("64748B")
    document.save(target)
    normalize_ooxml(target)


def email_bytes(renewal: Renewal) -> bytes:
    if renewal.action != "DRAFT_NOTIFICATION" or not renewal.primary_email or not renewal.primary_name:
        raise ValueError("只能为通过控制的账户生成邮件")
    message = EmailMessage(policy=SMTP)
    message["From"] = "me@ourcompany.example.com"
    message["To"] = renewal.primary_email
    message["Subject"] = f"Q1 Renewal Heads-Up — {renewal.account}"
    message["Date"] = "Tue, 24 Feb 2026 16:00:00 +0000"
    message["Message-ID"] = renewal.message_id
    message["X-EvoCowork-Delivery-Status"] = "DRAFT_NOT_SENT"
    message["X-EvoCowork-Source-Account-ID"] = renewal.account_id
    message["X-EvoCowork-Source-Contact-ID"] = renewal.primary_contact_id or ""
    message.set_content(
        f"Draft status: DRAFT_NOT_SENT. This message has not been sent.\n\n"
        f"Hello {renewal.primary_name},\n\n"
        f"We are preparing the Q1 renewal review for {renewal.account}. "
        f"The current contract value in our renewal tracker is {renewal.contract_text}, "
        f"with a renewal date of {renewal.renewal_date.isoformat()}.\n\n"
        "This is an offline heads-up draft for internal review before any customer communication.\n\n"
        "Best,\nRenewal Operations\n"
    )
    buffer = BytesIO()
    BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
    return buffer.getvalue()


def build_mbox(renewals: list[Renewal], target: Path) -> None:
    chunks: list[bytes] = []
    for renewal in renewals:
        if renewal.action != "DRAFT_NOTIFICATION":
            continue
        chunks.append(b"From me@ourcompany.example.com Tue Feb 24 16:00:00 2026\n")
        chunks.append(email_bytes(renewal).replace(b"\r\n", b"\n"))
        chunks.append(b"\n")
    target.write_bytes(b"".join(chunks))


def build_slack(renewals: list[Renewal], channel: dict, target: Path) -> None:
    eligible = [item for item in renewals if item.action == "DRAFT_NOTIFICATION"]
    total = sum((item.contract_amount for item in eligible), Decimal("0"))
    channel_record = channel["channels"][0]
    text = (
        "Q1 renewal notifications prepared as offline drafts: "
        + "; ".join(
            f"{item.account} ({item.contract_text}, renews {item.renewal_date.isoformat()})"
            for item in eligible
        )
        + f". Total prepared value: ${total:,.0f} across {len(eligible)} customer drafts. "
        "Customer messages have not been sent; this channel update has not been posted."
    )
    payload = {
        "task_id": "cowork-evo/automationbench-1202-q1-renewal-control-package",
        "channel": {"id": channel_record["id"], "name": channel_record["name"]},
        "message_id": "draft-renewals-1202",
        "status": "DRAFT_NOT_POSTED",
        "linked_message_ids": [item.message_id for item in eligible],
        "source_files": [
            "Renewal_Tracker.xlsx", "Renewal_Inbox.mbox",
            "Salesforce_Renewal_Contacts.json", "Renewals_Channel.json",
        ],
        "text": text,
    }
    target.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()

    renewals, accounts, contacts, channel, _messages = read_inputs(args.input_dir)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for existing in args.output_dir.iterdir():
        if existing.is_file() or existing.is_symlink():
            existing.unlink()
        elif existing.is_dir():
            shutil.rmtree(existing)

    build_workbook(renewals, accounts, contacts, args.output_dir / XLSX_NAME)
    build_docx(renewals, args.output_dir / DOCX_NAME)
    build_mbox(renewals, args.output_dir / MBOX_NAME)
    build_slack(renewals, channel, args.output_dir / SLACK_NAME)


if __name__ == "__main__":
    main()
