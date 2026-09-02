#!/usr/bin/env python3
"""AutomationBench 1202 的逐项定向负例回放。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Callable

from docx import Document
from openpyxl import load_workbook


XLSX = "Renewal_Decision_Register.xlsx"
DOCX = "Q1_Renewal_Control_Brief.docx"
MBOX = "Q1_Renewal_Communications.mbox"
SLACK = "Q1_Renewals_Channel_Draft.json"


def mutate_workbook(root: Path, callback: Callable) -> None:
    path = root / XLSX
    workbook = load_workbook(path)
    callback(workbook)
    workbook.save(path)


def mutate_document(root: Path, old: str, new: str) -> None:
    path = root / DOCX
    document = Document(path)
    changed = 0
    for paragraph in document.paragraphs:
        if old in paragraph.text:
            for run in paragraph.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
                    changed += 1
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if old in run.text:
                            run.text = run.text.replace(old, new)
                            changed += 1
    if not changed:
        raise RuntimeError(f"Word 负例未命中：{old}")
    document.save(path)


def mutate_bytes(root: Path, name: str, old: bytes, new: bytes, count: int = 1) -> None:
    path = root / name
    data = path.read_bytes()
    if old not in data:
        raise RuntimeError(f"{name} 负例未命中：{old!r}")
    path.write_bytes(data.replace(old, new, count))


def mutate_json(root: Path, callback: Callable) -> None:
    path = root / SLACK
    payload = json.loads(path.read_text(encoding="utf-8"))
    callback(payload)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def cases() -> dict[str, Callable[[Path], None]]:
    def missing_docx(root: Path) -> None:
        (root / DOCX).unlink()

    def extra_debug(root: Path) -> None:
        (root / "debug.txt").write_text("debug\n", encoding="utf-8")

    def source_account_changed(root: Path) -> None:
        mutate_workbook(root, lambda wb: setattr(wb["Source_Renewals"]["B2"], "value", "Vanguard Logistic"))

    def source_amount_changed(root: Path) -> None:
        mutate_workbook(root, lambda wb: setattr(wb["Source_Renewals"]["E3"], "value", "$124,000"))

    def source_hold_removed(root: Path) -> None:
        mutate_workbook(root, lambda wb: setattr(wb["Source_Renewals"]["G4"], "value", ""))

    def pinnacle_marked_hold(root: Path) -> None:
        mutate_workbook(root, lambda wb: setattr(wb["Authority_Control"]["G3"], "value", "HOLD_LEGAL"))

    def cobalt_marked_clear(root: Path) -> None:
        mutate_workbook(root, lambda wb: setattr(wb["Authority_Control"]["G4"], "value", "CLEAR"))

    def external_authorized(root: Path) -> None:
        mutate_workbook(root, lambda wb: setattr(wb["Authority_Control"]["C6"], "value", "AUTHORITATIVE_INTERNAL"))

    def secondary_contact_selected(root: Path) -> None:
        def edit(wb):
            wb["CRM_Contacts"]["G2"] = False
            wb["CRM_Contacts"]["G3"] = True
        mutate_workbook(root, edit)

    def recipient_secondary(root: Path) -> None:
        mutate_workbook(root, lambda wb: setattr(wb["Action_Register"]["I2"], "value", "tom.alvarez@vanguardlog.example.com"))

    def cobalt_drafted(root: Path) -> None:
        mutate_workbook(root, lambda wb: setattr(wb["Action_Register"]["J4"], "value", "DRAFT_NOTIFICATION"))

    def terraform_drafted(root: Path) -> None:
        mutate_workbook(root, lambda wb: setattr(wb["Action_Register"]["J5"], "value", "DRAFT_NOTIFICATION"))

    def decision_staticized(root: Path) -> None:
        mutate_workbook(root, lambda wb: setattr(wb["Action_Register"]["J2"], "value", "DRAFT_NOTIFICATION"))

    def summary_wrong_total(root: Path) -> None:
        mutate_workbook(root, lambda wb: setattr(wb["Delivery_Summary"]["C4"], "value", 226000))

    def chart_removed(root: Path) -> None:
        mutate_workbook(root, lambda wb: setattr(wb["Delivery_Summary"], "_charts", []))

    def tables_removed(root: Path) -> None:
        def edit(wb):
            for sheet in wb.worksheets:
                sheet.tables.clear()
        mutate_workbook(root, edit)

    def mbox_drop_message(root: Path) -> None:
        data = (root / MBOX).read_bytes()
        marker = b"\nFrom me@ourcompany.example.com Tue Feb 24 16:00:00 2026\n"
        if marker not in data:
            raise RuntimeError("MBOX 无第二封分隔符")
        (root / MBOX).write_bytes(data.split(marker, 1)[0] + b"\n")

    def mbox_secondary_recipient(root: Path) -> None:
        mutate_bytes(root, MBOX, b"sarah.chen@vanguardlog.example.com", b"tom.alvarez@vanguardlog.example.com")

    def mbox_wrong_amount(root: Path) -> None:
        mutate_bytes(root, MBOX, b"$142,000", b"$124,000")

    def mbox_false_sent(root: Path) -> None:
        mutate_bytes(root, MBOX, b"This message has not been sent.", b"This message has been sent successfully.")

    def mbox_message_id_conflict(root: Path) -> None:
        mutate_bytes(root, MBOX, b"renewal-ren2@evocowork.invalid", b"renewal-ren1@evocowork.invalid")

    def slack_wrong_total(root: Path) -> None:
        mutate_json(root, lambda payload: payload.__setitem__("text", payload["text"].replace("$227,000", "$226,000")))

    def slack_add_excluded(root: Path) -> None:
        mutate_json(root, lambda payload: payload.__setitem__("text", payload["text"] + " Cobalt Engineering also reviewed."))

    def slack_false_posted(root: Path) -> None:
        def edit(payload):
            payload["status"] = "POSTED"
            payload["text"] = payload["text"].replace("has not been posted", "has been posted")
        mutate_json(root, edit)

    def slack_message_id_conflict(root: Path) -> None:
        mutate_json(root, lambda payload: payload.__setitem__("linked_message_ids", ["<renewal-ren1@evocowork.invalid>"]))

    def word_wrong_total(root: Path) -> None:
        mutate_document(root, "$227,000", "$226,000")

    def word_external_authorized(root: Path) -> None:
        mutate_document(root, "External partner requests cannot add accounts", "External partner requests can add accounts")

    def word_false_sent(root: Path) -> None:
        mutate_document(root, "No email was sent", "Email has been sent successfully")

    return {
        "missing-docx": missing_docx,
        "extra-debug-file": extra_debug,
        "source-account-changed": source_account_changed,
        "source-amount-changed": source_amount_changed,
        "source-hold-removed": source_hold_removed,
        "pinnacle-marked-hold": pinnacle_marked_hold,
        "cobalt-marked-clear": cobalt_marked_clear,
        "external-authorized": external_authorized,
        "secondary-contact-selected": secondary_contact_selected,
        "action-recipient-secondary": recipient_secondary,
        "cobalt-drafted": cobalt_drafted,
        "terraform-drafted": terraform_drafted,
        "decision-staticized": decision_staticized,
        "summary-wrong-total": summary_wrong_total,
        "native-chart-removed": chart_removed,
        "native-tables-removed": tables_removed,
        "mbox-drop-message": mbox_drop_message,
        "mbox-secondary-recipient": mbox_secondary_recipient,
        "mbox-wrong-amount": mbox_wrong_amount,
        "mbox-false-sent-claim": mbox_false_sent,
        "mbox-message-id-conflict": mbox_message_id_conflict,
        "slack-wrong-total": slack_wrong_total,
        "slack-add-excluded": slack_add_excluded,
        "slack-false-posted": slack_false_posted,
        "slack-message-id-conflict": slack_message_id_conflict,
        "word-wrong-total": word_wrong_total,
        "word-external-authorized": word_external_authorized,
        "word-false-sent-claim": word_false_sent,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--golden", type=Path, default=Path(__file__).resolve().parents[1] / "solution" / "Golden")
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()

    judge_path = Path(__file__).with_name("judge.py")
    spec = importlib.util.spec_from_file_location("renewal_task_judge", judge_path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入本题 Judge")
    judge = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = judge
    spec.loader.exec_module(judge)

    results: list[dict] = []
    with tempfile.TemporaryDirectory(prefix="renewal-negative-cases-") as temporary:
        root = Path(temporary)
        for name, mutate in cases().items():
            case_dir = root / name
            shutil.copytree(args.golden, case_dir)
            mutate(case_dir)
            result = judge.evaluate(case_dir)
            results.append({
                "case": name,
                "reward": result["reward"],
                "passed": result["passed"],
                "hurdle_failures": result.get("hurdle_failures", []),
                "penalties": [item["id"] for item in result.get("penalties", [])],
            })

    failed = [item for item in results if item["passed"] or float(item["reward"]) >= 0.8]
    report = {
        "total": len(results),
        "rejected": len(results) - len(failed),
        "maximum_reward": max(float(item["reward"]) for item in results),
        "cases": results,
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if failed:
        raise SystemExit(f"存在误接受负例：{[item['case'] for item in failed]}")


if __name__ == "__main__":
    main()
