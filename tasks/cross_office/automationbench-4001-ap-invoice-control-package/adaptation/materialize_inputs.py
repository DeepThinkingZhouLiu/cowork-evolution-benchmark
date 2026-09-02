#!/usr/bin/env python3
"""将 AutomationBench 4001 的公开初始状态逐字段静态化为本题输入。

这不是通用制题器：脚本只识别 4001，并对这道题的邮件、表格和
Slack 字段做完整性断言。保留脚本是为了让输入的来源转换可审计。
"""

from __future__ import annotations

import importlib.util
import json
import shutil
import sys
import types
import zipfile
from datetime import datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
UPSTREAM = Path(__file__).with_name("upstream-finance-tasks.py")
SNAPSHOT = Path(__file__).with_name("upstream-task-4001.json")
DATA_DIR = ROOT / "data" / "input_files"
ENV_DIR = ROOT / "environment" / "input_files"
FIXED_ZIP_TIME = (2026, 2, 3, 9, 0, 0)


def load_upstream_task() -> dict:
    datasets = types.ModuleType("datasets")
    datasets.Dataset = object
    sys.modules["datasets"] = datasets
    for name in ("automationbench", "automationbench.domains", "automationbench.domains.finance"):
        sys.modules[name] = types.ModuleType(name)
    noise = types.ModuleType("automationbench.domains.finance._noise")
    noise.apply_noise = lambda value: value
    sys.modules[noise.__name__] = noise

    spec = importlib.util.spec_from_file_location("automationbench_finance_tasks_frozen", UPSTREAM)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入冻结的 AutomationBench 源文件")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    task = module.get_fin_invoice_email_extract_task()
    if task.get("example_id") != 4001 or task.get("task") != "finance.invoice_email_extract":
        raise ValueError("冻结函数不再是 AutomationBench 4001")
    return task


def assert_task_shape(task: dict) -> None:
    state = task["info"]["initial_state"]
    messages = state["gmail"]["messages"]
    worksheets = state["google_sheets"]["spreadsheets"][0]["worksheets"]
    slack_messages = state["slack"]["messages"]
    assert [message["id"] for message in messages] == [
        "msg_inv_001",
        "msg_inv_002",
        "msg_inv_003",
        "msg_inv_004",
        "msg_ap_policy",
        "msg_vp_override",
        "msg_ext_override",
    ]
    assert [sheet["id"] for sheet in worksheets] == ["ws_pending", "ws_blocked"]
    assert len(worksheets[0]["rows"]) == 0
    assert len(worksheets[1]["rows"]) == 3
    assert len(slack_messages) == 1 and slack_messages[0]["id"] == "slk_gl_corr"
    assert len(task["info"]["assertions"]) == 9


def message_bytes(record: dict) -> bytes:
    message = EmailMessage(policy=SMTP)
    message["From"] = record["from_"]
    message["To"] = ", ".join(record["to"])
    message["Subject"] = record["subject"]
    parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
    message["Date"] = parsed.strftime("%a, %d %b %Y %H:%M:%S +0000")
    message["Message-ID"] = f"<{record['id']}@automationbench.invalid>"
    message["X-AutomationBench-Message-ID"] = record["id"]
    message["X-AutomationBench-Thread-ID"] = record["thread_id"]
    message["X-AutomationBench-Labels"] = ",".join(record["label_ids"])
    message["X-AutomationBench-Is-Read"] = str(record["is_read"]).lower()
    message.set_content(record["body_plain"], charset="utf-8")
    buffer = BytesIO()
    BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
    return buffer.getvalue()


def write_mbox(messages: list[dict], target: Path) -> None:
    chunks: list[bytes] = []
    for record in messages:
        sender = record["from_"]
        parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
        chunks.append(f"From {sender} {parsed.strftime('%a %b %d %H:%M:%S %Y')}\n".encode("ascii"))
        chunks.append(message_bytes(record).replace(b"\r\n", b"\n"))
        chunks.append(b"\n")
    target.write_bytes(b"".join(chunks))


def add_table(sheet, name: str, reference: str) -> None:
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def normalize_xlsx(path: Path) -> None:
    temporary = path.with_suffix(".normalized.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        temporary, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as destination:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            destination.writestr(info, source.read(name))
    temporary.replace(path)


def write_tracker(state: dict, target: Path) -> None:
    worksheets = state["google_sheets"]["spreadsheets"][0]["worksheets"]
    pending, blocked = worksheets
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.creator = "AutomationBench 4001 public fixture"
    workbook.title = "Invoice Tracker"
    workbook.subject = "Static Office representation of Google Sheets initial state"
    workbook.description = "Materialized without adding business records beyond AutomationBench 4001."

    pending_sheet = workbook.create_sheet(pending["title"])
    pending_sheet.append(pending["headers"])
    add_table(pending_sheet, "PendingInvoicesInput", "A1:D1")

    blocked_sheet = workbook.create_sheet(blocked["title"])
    blocked_headers = ["Vendor Name", "Reason", "Status", "Notes"]
    blocked_sheet.append(blocked_headers)
    for row in blocked["rows"]:
        cells = row["cells"]
        blocked_sheet.append([cells.get(header, "") for header in blocked_headers])
    add_table(blocked_sheet, "BlockedVendorsInput", f"A1:D{blocked_sheet.max_row}")

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center")
        widths = {"A": 34, "B": 48, "C": 20, "D": 68}
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_area = sheet.dimensions

    workbook.save(target)
    normalize_xlsx(target)


def main() -> None:
    task = load_upstream_task()
    assert_task_shape(task)
    state = task["info"]["initial_state"]
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    ENV_DIR.mkdir(parents=True, exist_ok=True)

    SNAPSHOT.write_text(json.dumps(task, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_mbox(state["gmail"]["messages"], DATA_DIR / "AP_Inbox.mbox")
    write_tracker(state, DATA_DIR / "Invoice_Tracker.xlsx")
    slack_payload = {
        "meta": state["meta"],
        "channels": state["slack"]["channels"],
        "messages": state["slack"]["messages"],
    }
    (DATA_DIR / "Accounts_Payable_Slack.json").write_text(
        json.dumps(slack_payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    for source in sorted(DATA_DIR.iterdir()):
        if source.is_file():
            shutil.copy2(source, ENV_DIR / source.name)


if __name__ == "__main__":
    main()
