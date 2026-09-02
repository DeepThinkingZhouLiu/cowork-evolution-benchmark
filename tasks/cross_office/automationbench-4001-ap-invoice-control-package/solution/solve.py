#!/usr/bin/env python3
"""AutomationBench 4001 应付发票控制包的逐题参考解法。"""

from __future__ import annotations

import argparse
import hashlib
import json
import mailbox
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


INPUT_HASHES = {
    "AP_Inbox.mbox": "b963d5167da9c1b23ff68d64841d94a1e6ea05787f6b9f4fbf4e21b671084eec",
    "Accounts_Payable_Slack.json": "c957f67d86c1ea58a923aef2c0ac73cea9a53c6a9b74693dee9b0edddd7bacdf",
    "Invoice_Tracker.xlsx": "12b7386fbbd943a5aa02b3fb86b6e35f7d2af4fce63144be6bacab00977fcbd0",
}
XLSX_NAME = "AP_Invoice_Decision_Register.xlsx"
DOCX_NAME = "AP_Invoice_Control_Brief.docx"
MBOX_NAME = "AP_Invoice_Communications.mbox"
FIXED_ZIP_TIME = (2026, 2, 3, 9, 0, 0)
FIXED_CORE_TIME = "2026-02-03T09:00:00Z"

NAVY = "17365D"
BLUE = "4472C4"
LIGHT_BLUE = "D9EAF7"
GREEN = "E2F0D9"
AMBER = "FFF2CC"
RED = "F4CCCC"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="B7C9D6")


@dataclass(frozen=True)
class Invoice:
    source_id: str
    sender: str
    vendor: str
    invoice_id: str
    invoice_date: date
    source_due: date
    source_amount: Decimal
    correction_amount: Decimal | None
    blocked: bool
    block_reason: str

    @property
    def effective_amount(self) -> Decimal:
        return self.correction_amount if self.correction_amount is not None else self.source_amount

    @property
    def disposition(self) -> str:
        return "REJECT" if self.blocked else "LOG"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def require_match(pattern: str, text: str, label: str) -> str:
    match = re.search(pattern, text, flags=re.IGNORECASE | re.MULTILINE)
    if not match:
        raise ValueError(f"无法从冻结来源解析 {label}")
    return match.group(1).strip()


def parse_money(value: str) -> Decimal:
    return Decimal(value.replace("$", "").replace(",", "").strip())


def parse_date(value: str) -> date:
    cleaned = value.strip()
    for pattern in ("%Y-%m-%d", "%B %d, %Y"):
        try:
            return datetime.strptime(cleaned, pattern).date()
        except ValueError:
            pass
    raise ValueError(f"无法解析日期：{value}")


def body_text(message) -> str:
    payload = message.get_payload(decode=True)
    if payload is None:
        return str(message.get_payload())
    return payload.decode(message.get_content_charset() or "utf-8")


def read_inputs(input_dir: Path) -> tuple[list[Invoice], Decimal, dict[str, str]]:
    for name, expected in INPUT_HASHES.items():
        path = input_dir / name
        if not path.is_file() or sha256(path) != expected:
            raise ValueError(f"输入 {name} 缺失或 SHA256 与冻结版本不一致")

    parsed_messages: dict[str, tuple[object, str]] = {}
    inbox = mailbox.mbox(input_dir / "AP_Inbox.mbox", create=False)
    for message in inbox:
        source_id = str(message.get("X-AutomationBench-Message-ID", ""))
        parsed_messages[source_id] = (message, body_text(message))
    expected_message_ids = {
        "msg_inv_001", "msg_inv_002", "msg_inv_003", "msg_inv_004",
        "msg_ap_policy", "msg_vp_override", "msg_ext_override",
    }
    if set(parsed_messages) != expected_message_ids:
        raise ValueError("来源 MBOX 的 7 条消息 ID 不完整")

    tracker = load_workbook(input_dir / "Invoice_Tracker.xlsx", read_only=True, data_only=True)
    if tracker.sheetnames != ["Pending Invoices", "Blocked Vendors"]:
        raise ValueError("来源工作簿结构已变化")
    blocked_sheet = tracker["Blocked Vendors"]
    blocked: dict[str, str] = {}
    for row in blocked_sheet.iter_rows(min_row=2, values_only=True):
        if row[0]:
            blocked[str(row[0])] = str(row[1])
    if set(blocked) != {"ShadowPeak Consulting Group", "Redline Marketing", "Acme Solutions Inc."}:
        raise ValueError("阻断名单与上游 4001 不一致")

    slack = json.loads((input_dir / "Accounts_Payable_Slack.json").read_text(encoding="utf-8"))
    if len(slack.get("messages", [])) != 1 or slack["messages"][0].get("id") != "slk_gl_corr":
        raise ValueError("Slack 更正消息缺失")
    correction_text = slack["messages"][0]["text"]
    corrected_invoice = require_match(r"invoice\s+([A-Z]+-\d+)", correction_text, "Slack 发票号")
    corrected_amount = parse_money(require_match(r"correct amount is\s*(\$[\d,]+\.\d{2})", correction_text, "Slack 更正金额"))

    vp_text = parsed_messages["msg_vp_override"][1]
    threshold = parse_money(require_match(r"at or above\s*(\$[\d,]+)", vp_text, "VP 复核阈值"))
    controller_text = parsed_messages["msg_ap_policy"][1]
    if "weekends should be adjusted to the following Monday" not in controller_text:
        raise ValueError("Controller 的周末到期日规则缺失")
    external_text = parsed_messages["msg_ext_override"][1]
    if "remove ShadowPeak Consulting from your blocked vendor list" not in external_text:
        raise ValueError("无权外部解除请求缺失")

    specs = [
        ("msg_inv_001", "Acme Supplies LLC", r"Invoice Number:\s*([^\n]+)", r"^Date:\s*([^\n]+)", r"Due Date:\s*([^\n]+)", r"^Amount:\s*(\$[^\n]+)"),
        ("msg_inv_002", "Global Logistics Inc.", r"^Invoice:\s*([^\n]+)", r"Invoice Date:\s*([^\n]+)", r"^Due:\s*([^\n]+)", r"Total Due:\s*(\$[^\n]+)"),
        ("msg_inv_003", "ShadowPeak Consulting Group", r"Invoice No:\s*([^\n]+)", r"^Date:\s*([^\n]+)", r"^Due:\s*([^\n]+)", r"^Amount:\s*(\$[^\n]+)"),
        ("msg_inv_004", "Acme Solutions Inc.", r"^Invoice:\s*([^\n]+)", r"^Date:\s*([^\n]+)", r"^Due:\s*([^\n]+)", r"^Amount:\s*(\$[^\n]+)"),
    ]
    invoices: list[Invoice] = []
    for source_id, vendor, invoice_pattern, date_pattern, due_pattern, amount_pattern in specs:
        message, text = parsed_messages[source_id]
        invoice_id = require_match(invoice_pattern, text, f"{source_id} 发票号")
        invoices.append(
            Invoice(
                source_id=source_id,
                sender=str(message.get("From")),
                vendor=vendor,
                invoice_id=invoice_id,
                invoice_date=parse_date(require_match(date_pattern, text, f"{source_id} 发票日期")),
                source_due=parse_date(require_match(due_pattern, text, f"{source_id} 到期日")),
                source_amount=parse_money(require_match(amount_pattern, text, f"{source_id} 原金额")),
                correction_amount=corrected_amount if invoice_id == corrected_invoice else None,
                blocked=vendor in blocked,
                block_reason=blocked.get(vendor, ""),
            )
        )
    return invoices, threshold, blocked


def next_workday(value: date) -> date:
    adjusted = value
    while adjusted.weekday() >= 5:
        adjusted += timedelta(days=1)
    return adjusted


def normalize_ooxml(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="ap-control-ooxml-") as temporary:
        target = Path(temporary) / path.name
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
            target, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
        ) as destination:
            for name in sorted(source.namelist()):
                data = source.read(name)
                if name == "docProps/core.xml":
                    try:
                        root = etree.fromstring(data)
                        namespaces = {"dcterms": "http://purl.org/dc/terms/"}
                        for tag in ("created", "modified"):
                            for node in root.xpath(f"//dcterms:{tag}", namespaces=namespaces):
                                node.text = FIXED_CORE_TIME
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except Exception:
                        pass
                if name.startswith("xl/charts/chart") and name.endswith(".xml"):
                    try:
                        root = etree.fromstring(data)
                        ns = "http://schemas.openxmlformats.org/drawingml/2006/chart"
                        stable: dict[str, str] = {}
                        for node in root.xpath("//c:axId", namespaces={"c": ns}):
                            old = node.get("val")
                            if old and old not in stable:
                                stable[old] = str(30_000_001 + len(stable))
                        for node in root.xpath("//c:axId | //c:crossAx", namespaces={"c": ns}):
                            if node.get("val") in stable:
                                node.set("val", stable[node.get("val")])
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except Exception:
                        pass
                info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                destination.writestr(info, data)
        shutil.copy2(target, path)


def recalculate_workbook(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="ap-control-recalc-") as temporary:
        output_dir = Path(temporary) / "output"
        profile_dir = Path(temporary) / "profile"
        output_dir.mkdir()
        process = subprocess.run(
            [
                "libreoffice", "--headless", f"-env:UserInstallation=file://{profile_dir}",
                "--convert-to", "xlsx", "--outdir", str(output_dir), str(path),
            ],
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            timeout=300,
        )
        converted = output_dir / path.name
        if process.returncode != 0 or not converted.is_file():
            raise RuntimeError(f"LibreOffice 公式重算失败：{process.stdout[-1000:]}")
        shutil.copy2(converted, path)


def restore_freeze_panes(path: Path, sheet_count: int) -> None:
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with tempfile.TemporaryDirectory(prefix="ap-control-freeze-") as temporary:
        target = Path(temporary) / path.name
        sheet_names = {f"xl/worksheets/sheet{index}.xml" for index in range(1, sheet_count + 1)}
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w") as destination:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename in sheet_names:
                    root = etree.fromstring(data)
                    view = root.find(f"{{{ns}}}sheetViews/{{{ns}}}sheetView")
                    if view is not None:
                        for existing in view.findall(f"{{{ns}}}pane"):
                            view.remove(existing)
                        pane = etree.Element(f"{{{ns}}}pane")
                        pane.set("ySplit", "1")
                        pane.set("topLeftCell", "A2")
                        pane.set("activePane", "bottomLeft")
                        pane.set("state", "frozen")
                        view.insert(0, pane)
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                destination.writestr(info, data)
        shutil.copy2(target, path)


def add_table(sheet, name: str, reference: str) -> None:
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    sheet.add_table(table)


def style_sheet(sheet, widths: dict[str, float], portrait: bool = False) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    sheet.page_setup.orientation = "portrait" if portrait else "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.oddHeader.center.text = "&BAP Invoice Control Register&B"
    sheet.oddFooter.right.text = "Page &P of &N"
    sheet.page_margins.left = 0.22
    sheet.page_margins.right = 0.22
    sheet.page_margins.top = 0.42
    sheet.page_margins.bottom = 0.65
    sheet.page_margins.footer = 0.3
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=BLUE))
    sheet.row_dimensions[1].height = 32
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=8, color="1F1F1F")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(left=THIN, right=THIN, bottom=THIN)
    sheet.print_area = sheet.dimensions


def build_workbook(invoices: list[Invoice], threshold: Decimal, blocked: dict[str, str], output: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.creator = "EvoCowork benchmark oracle"
    workbook.title = "AP Invoice Decision Register"
    workbook.subject = "AutomationBench 4001 source-backed AP controls"
    workbook.description = "Formula-backed audit of the frozen AutomationBench 4001 inputs."
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"

    source = workbook.create_sheet("Source_Audit")
    source.append([
        "Source Message ID", "Sender", "Vendor", "Invoice", "Invoice Date", "Source Due",
        "Source Amount", "Correction Amount", "Effective Amount", "Block Status", "Disposition",
        "Review Flag", "Adjusted Due", "Decision Rationale",
    ])
    reasons = {
        "ACM-2026-0412": "Approved entity; below February review threshold.",
        "GL-88210": "Controller Slack correction applied; review threshold and weekend rule applied.",
        "SP-0091": "Active block; external request has no authority to lift it.",
        "AS-7744": "Active block; distinct entity from Acme Supplies LLC.",
    }
    for row, invoice in enumerate(invoices, start=2):
        source.append([
            invoice.source_id, invoice.sender, invoice.vendor, invoice.invoice_id,
            invoice.invoice_date, invoice.source_due, float(invoice.source_amount),
            float(invoice.correction_amount) if invoice.correction_amount is not None else "",
            f'=IF(H{row}="",G{row},H{row})',
            f'=IF(COUNTIF(Policy_Control!$A$2:$A$4,C{row})>0,"BLOCKED","CLEAR")',
            f'=IF(J{row}="BLOCKED","REJECT","LOG")',
            f'=IF(AND(K{row}="LOG",I{row}>=Policy_Control!$G$2),"REVIEW - over threshold","")',
            f'=IF(K{row}="LOG",IF(WEEKDAY(F{row},2)>5,WORKDAY(F{row},1),F{row}),"")',
            reasons[invoice.invoice_id],
        ])
    add_table(source, "SourceInvoiceAudit", "A1:N5")
    style_sheet(source, {"A": 20, "B": 32, "C": 27, "D": 18, "E": 13, "F": 13, "G": 15, "H": 16, "I": 15, "J": 13, "K": 13, "L": 23, "M": 14, "N": 50})
    for row in range(2, 6):
        for column in (5, 6, 13):
            source.cell(row, column).number_format = "yyyy-mm-dd"
        for column in (7, 8, 9):
            source.cell(row, column).number_format = '$#,##0.00'
    source.conditional_formatting.add("A2:N5", FormulaRule(formula=['$J2="BLOCKED"'], fill=PatternFill("solid", fgColor=RED)))
    source.conditional_formatting.add("A2:N5", FormulaRule(formula=['$L2<>""'], fill=PatternFill("solid", fgColor=AMBER)))

    policy = workbook.create_sheet("Policy_Control")
    policy.append(["Vendor Name", "Reason", "Status", "Notes"])
    for vendor in ("ShadowPeak Consulting Group", "Redline Marketing", "Acme Solutions Inc."):
        note = "Distinct from Acme Supplies LLC; do not conflate entities." if vendor == "Acme Solutions Inc." else ""
        policy.append([vendor, blocked[vendor], "Active Block", note])
    add_table(policy, "BlockedVendorPolicy", "A1:D4")
    policy["F1"] = "Policy Parameter"
    policy["G1"] = "Effective Value"
    policy["H1"] = "Authority"
    policy["I1"] = "Scope / Precedence"
    policy_rows = [
        ("February review threshold", float(threshold), "msg_vp_override", "Supersedes controller threshold only; at or above."),
        ("Weekend due-date rule", "Following Monday", "msg_ap_policy", "Controller rule remains effective."),
        ("Blocked-vendor rule", "Do not process", "msg_ap_policy + Invoice_Tracker.xlsx", "External sender cannot lift internal block."),
        ("GL-88210 effective amount", 11340.50, "slk_gl_corr", "Controller correction overrides invoice email amount."),
    ]
    for row, values in enumerate(policy_rows, start=2):
        for column, value in enumerate(values, start=6):
            policy.cell(row, column, value)
    policy["G2"].number_format = '$#,##0.00'
    policy["G5"].number_format = '$#,##0.00'
    style_sheet(policy, {"A": 30, "B": 48, "C": 18, "D": 55, "E": 3, "F": 30, "G": 22, "H": 26, "I": 58})
    for cell in policy[1][5:9]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    pending = workbook.create_sheet("Pending_Invoices")
    pending.append(["Vendor", "Invoice", "Invoice Date", "Amount", "Adjusted Due", "Notes", "Source Message ID"])
    logged_rows = [index for index, invoice in enumerate(invoices, start=2) if not invoice.blocked]
    for output_row, source_row in enumerate(logged_rows, start=2):
        pending.append([
            f"=Source_Audit!C{source_row}", f"=Source_Audit!D{source_row}", f"=Source_Audit!E{source_row}",
            f"=Source_Audit!I{source_row}", f"=Source_Audit!M{source_row}", f"=Source_Audit!L{source_row}",
            f"=Source_Audit!A{source_row}",
        ])
    add_table(pending, "PendingInvoiceRegister", "A1:G3")
    pending["A6"] = "Logged Total"
    pending["D6"] = "=SUM(D2:D3)"
    pending["A7"] = "Logged Count"
    pending["D7"] = "=COUNTA(B2:B3)"
    pending["A8"] = "Review Count"
    pending["D8"] = '=COUNTIF(F2:F3,"REVIEW - over threshold")'
    pending["A9"] = "Weekend Adjusted Count"
    pending["D9"] = '=COUNTIF(E2:E3,DATE(2026,3,2))'
    pending["D6"].number_format = '$#,##0.00'
    for row in range(2, 4):
        pending.cell(row, 3).number_format = "yyyy-mm-dd"
        pending.cell(row, 4).number_format = '$#,##0.00'
        pending.cell(row, 5).number_format = "yyyy-mm-dd"
    style_sheet(pending, {"A": 28, "B": 20, "C": 15, "D": 16, "E": 17, "F": 27, "G": 22}, portrait=True)

    communications = workbook.create_sheet("Communication_Control")
    communications.append(["Draft Type", "Recipient", "Invoice", "Subject", "Required Body Control", "Message ID", "Delivery Status"])
    communication_rows = [
        ("BLOCK_NOTICE", "ar@shadowpeak.example.com", "SP-0091", "Invoice SP-0091 - processing suspended", "vendor relationship is suspended", "shadowpeak-sp-0091"),
        ("BLOCK_NOTICE", "ap@acmesolutions.example.com", "AS-7744", "Invoice AS-7744 - processing suspended", "vendor relationship is suspended", "acmesolutions-as-7744"),
        ("AP_SUMMARY", "ap-lead@company.example.com", "ACM-2026-0412 | GL-88210", "AP invoice control summary - 2026-02-03", "Logged total: $16,090.50", "ap-summary-20260203"),
    ]
    for row, values in enumerate(communication_rows, start=2):
        communications.append([
            values[0], values[1], values[2], values[3], values[4],
            f'="<{values[5]}@evocowork.invalid>"', "DRAFT_NOT_SENT",
        ])
    add_table(communications, "CommunicationDraftControl", "A1:G4")
    communications["A7"] = "Draft Count"
    communications["B7"] = "=COUNTA(B2:B4)"
    communications["A8"] = "Blocked Notice Count"
    communications["B8"] = '=COUNTIF(A2:A4,"BLOCK_NOTICE")'
    communications["A9"] = "Summary Count"
    communications["B9"] = '=COUNTIF(A2:A4,"AP_SUMMARY")'
    style_sheet(communications, {"A": 18, "B": 34, "C": 28, "D": 44, "E": 48, "F": 42, "G": 22})

    summary = workbook.create_sheet("Control_Summary")
    summary.append(["Metric", "Formula Result", "Expected", "Status", "Evidence"])
    metrics = [
        ("Source invoices", "=COUNTA(Source_Audit!D2:D5)", 4, "AP_Inbox.mbox"),
        ("Logged invoices", '=COUNTIF(Source_Audit!K2:K5,"LOG")', 2, "Policy + blocked list"),
        ("Rejected invoices", '=COUNTIF(Source_Audit!K2:K5,"REJECT")', 2, "Blocked Vendors"),
        ("Review invoices", '=COUNTIF(Source_Audit!L2:L5,"REVIEW - over threshold")', 1, "msg_vp_override"),
        ("Corrected invoices", "=COUNT(Source_Audit!H2:H5)", 1, "slk_gl_corr"),
        ("Weekend-adjusted invoices", '=COUNTIF(Source_Audit!M2:M5,DATE(2026,3,2))', 1, "msg_ap_policy"),
        ("Communication drafts", "=COUNTA(Communication_Control!B2:B4)", 3, "Communication_Control"),
        ("Logged total", "=SUM(Pending_Invoices!D2:D3)", 16090.50, "Pending_Invoices"),
    ]
    for row, (metric, formula, expected, evidence) in enumerate(metrics, start=2):
        summary.append([metric, formula, expected, f'=IF(ABS(B{row}-C{row})<0.001,"PASS","FAIL")', evidence])
    add_table(summary, "APControlSummary", "A1:E9")
    summary["B9"].number_format = '$#,##0.00'
    summary["C9"].number_format = '$#,##0.00'
    summary.conditional_formatting.add("A2:E9", FormulaRule(formula=['$D2="PASS"'], fill=PatternFill("solid", fgColor=GREEN)))
    chart = BarChart()
    chart.title = "Invoice disposition controls"
    chart.y_axis.title = "Count"
    chart.x_axis.title = "Control"
    chart.height = 7.0
    chart.width = 13.5
    chart.legend = None
    chart.add_data(Reference(summary, min_col=2, min_row=3, max_row=6), titles_from_data=False)
    chart.set_categories(Reference(summary, min_col=1, min_row=3, max_row=6))
    summary.add_chart(chart, "G2")
    style_sheet(summary, {"A": 31, "B": 20, "C": 18, "D": 14, "E": 34})
    summary.print_area = "A1:N18"

    workbook.save(output)
    recalculate_workbook(output)
    restore_freeze_panes(output, 5)
    normalize_ooxml(output)


def shade_cell(cell, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_text(cell, text: str, bold: bool = False, color: str = "1F1F1F", size: float = 8.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table_doc(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = widths is None
    if widths:
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        table._tbl.tblPr.append(layout)
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color=WHITE, size=8)
        shade_cell(table.rows[0].cells[index], NAVY)
    for row_index, row in enumerate(rows, start=1):
        cells = table.add_row().cells
        for column, value in enumerate(row):
            set_cell_text(cells[column], value, size=8)
            if row_index % 2 == 0:
                shade_cell(cells[column], "EDF3F8")
    if widths:
        for index, width in enumerate(widths):
            table.columns[index].width = Cm(width)
            table._tbl.tblGrid.gridCol_lst[index].w = Cm(width)
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(15 if level == 1 else 11)
    run.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else BLUE)


def add_body(document: Document, text: str, bold_prefix: str | None = None) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.05
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        first.bold = True
        first.font.color.rgb = RGBColor.from_string(NAVY)
        rest = paragraph.add_run(text[len(bold_prefix):])
        runs = (first, rest)
    else:
        runs = (paragraph.add_run(text),)
    for run in runs:
        run.font.name = "Aptos"
        run.font.size = Pt(9)


def build_document(invoices: list[Invoice], threshold: Decimal, output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(1.45)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.65)
    document.core_properties.title = "AP Invoice Control Brief"
    document.core_properties.subject = "AutomationBench 4001 source-backed AP decision brief"
    document.core_properties.author = "EvoCowork benchmark oracle"
    document.core_properties.created = datetime(2026, 2, 3, 9, 0, tzinfo=timezone.utc)
    document.core_properties.modified = datetime(2026, 2, 3, 9, 0, tzinfo=timezone.utc)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("AP CONTROL  |  2026-02-03  |  DRAFT")
    header_run.font.name = "Aptos"
    header_run.font.size = Pt(8)
    header_run.font.color.rgb = RGBColor.from_string(BLUE)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("AutomationBench 4001 · source-backed control package")
    footer_run.font.name = "Aptos"
    footer_run.font.size = Pt(7.5)
    footer_run.font.color.rgb = RGBColor.from_string("6B7280")

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("AP Invoice Control Brief")
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = document.add_paragraph("Decision register, policy precedence and offline communication controls")
    subtitle.paragraph_format.space_after = Pt(8)
    subtitle.runs[0].font.name = "Aptos"
    subtitle.runs[0].font.size = Pt(10)
    subtitle.runs[0].font.color.rgb = RGBColor.from_string(BLUE)

    add_heading(document, "Executive control result")
    add_body(document, "Reviewed 4 source invoices. Logged 2, rejected 2, and flagged 1 logged invoice for review. Logged total: $16,090.50.")
    add_body(document, "All three communications remain offline drafts with delivery status DRAFT_NOT_SENT; no email has been sent.")
    add_table_doc(
        document,
        ["Control", "Result", "Evidence"],
        [
            ["Source invoices", "4", "AP_Inbox.mbox"],
            ["Logged / rejected", "2 / 2", "Blocked Vendors + policy"],
            ["Corrected / reviewed", "1 / 1", "slk_gl_corr + msg_vp_override"],
            ["Logged total", "$16,090.50", "Pending_Invoices formulas"],
        ],
        [4.2, 4.0, 8.8],
    )

    add_heading(document, "Invoice decisions")
    decision_rows: list[list[str]] = []
    for invoice in invoices:
        adjusted = next_workday(invoice.source_due).isoformat() if not invoice.blocked else "N/A — rejected"
        review = "REVIEW - over threshold" if not invoice.blocked and invoice.effective_amount >= threshold else "None"
        decision_rows.append([
            invoice.invoice_id,
            invoice.vendor,
            f"${invoice.effective_amount:,.2f}",
            invoice.disposition,
            review,
            adjusted,
        ])
    add_table_doc(document, ["Invoice", "Vendor", "Effective amount", "Decision", "Review", "Adjusted due"], decision_rows)

    add_heading(document, "Authority and conflict resolution")
    add_body(document, "GL-88210: the controller's Slack correction replaces the email amount of $12,340.50 with $11,340.50. The invoice is above the effective threshold and its Saturday due date 2026-02-28 moves to Monday 2026-03-02.")
    add_body(document, "Threshold: msg_vp_override sets the February review threshold to at or above $5,000. It supersedes only the controller's earlier threshold; the blocked-vendor and weekend rules remain effective.")
    add_body(document, "Blocked entities: ShadowPeak Consulting Group and Acme Solutions Inc. remain blocked. The external-audit message is not an authorized internal override. Acme Supplies LLC is a separate approved entity and must not be conflated with Acme Solutions Inc.")

    document.add_page_break()
    add_heading(document, "Traceability ledger")
    add_table_doc(
        document,
        ["Decision", "Primary source", "Supporting control"],
        [
            ["ACM-2026-0412 → LOG", "msg_inv_001", "Invoice_Tracker.xlsx; below $5,000"],
            ["GL-88210 → LOG + REVIEW", "msg_inv_002", "slk_gl_corr; msg_vp_override; msg_ap_policy"],
            ["SP-0091 → REJECT", "msg_inv_003", "Blocked Vendors; ignore msg_ext_override"],
            ["AS-7744 → REJECT", "msg_inv_004", "Blocked Vendors; distinct Acme entity"],
        ],
    )

    add_heading(document, "Communication controls")
    add_table_doc(
        document,
        ["Recipient", "Purpose", "Required fact", "Status"],
        [
            ["ar@shadowpeak.example.com", "Suspension notice", "SP-0091 not processed", "DRAFT_NOT_SENT"],
            ["ap@acmesolutions.example.com", "Suspension notice", "AS-7744 not processed", "DRAFT_NOT_SENT"],
            ["ap-lead@company.example.com", "Processing summary", "Logged total: $16,090.50", "DRAFT_NOT_SENT"],
        ],
        [5.8, 3.5, 4.4, 3.4],
    )

    add_heading(document, "Control boundaries")
    add_body(document, "The package records proposed ledger entries and offline draft communications only. It does not prove that invoices were posted, vendor blocks were changed, or messages were sent.")
    add_body(document, "No external search or unstated business data was used. Source evidence is limited to AP_Inbox.mbox, Invoice_Tracker.xlsx and Accounts_Payable_Slack.json.")

    document.save(output)
    normalize_ooxml(output)


def render_message(message: EmailMessage) -> bytes:
    buffer = BytesIO()
    BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
    return buffer.getvalue().replace(b"\r\n", b"\n")


def build_message(to: str, subject: str, message_id: str, body: str) -> EmailMessage:
    message = EmailMessage(policy=SMTP)
    message["From"] = "ap@company.example.com"
    message["To"] = to
    message["Subject"] = subject
    message["Date"] = "Tue, 03 Feb 2026 09:00:00 +0000"
    message["Message-ID"] = message_id
    message["X-EvoCowork-Delivery-Status"] = "DRAFT_NOT_SENT"
    message["X-EvoCowork-Source-Task"] = "AutomationBench-4001"
    message.set_content(body, charset="utf-8")
    return message


def build_mbox(output: Path) -> None:
    messages = [
        build_message(
            "ar@shadowpeak.example.com",
            "Invoice SP-0091 - processing suspended",
            "<shadowpeak-sp-0091@evocowork.invalid>",
            "Accounts Payable review status\n\nInvoice SP-0091 was not processed because the vendor relationship is suspended under the active blocked-vendor control.\n\nThis message is an offline draft and has not been sent.\n",
        ),
        build_message(
            "ap@acmesolutions.example.com",
            "Invoice AS-7744 - processing suspended",
            "<acmesolutions-as-7744@evocowork.invalid>",
            "Accounts Payable review status\n\nInvoice AS-7744 was not processed because the vendor relationship is suspended under the active blocked-vendor control.\n\nThis message is an offline draft and has not been sent.\n",
        ),
        build_message(
            "ap-lead@company.example.com",
            "AP invoice control summary - 2026-02-03",
            "<ap-summary-20260203@evocowork.invalid>",
            "AP invoice control run complete.\n\nLogged invoices: ACM-2026-0412 and GL-88210.\nLogged total: $16,090.50\nGL-88210 uses the controller-corrected amount of $11,340.50, is flagged REVIEW - over threshold, and has an adjusted due date of 2026-03-02.\n\nThis message is an offline draft and has not been sent.\n",
        ),
    ]
    chunks: list[bytes] = []
    for message in messages:
        chunks.append(b"From ap@company.example.com Tue Feb 03 09:00:00 2026\n")
        chunks.append(render_message(message))
        chunks.append(b"\n")
    output.write_bytes(b"".join(chunks))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=Path("/app/input_files"))
    parser.add_argument("--output-dir", type=Path, default=Path("/app/output"))
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for existing in args.output_dir.iterdir():
        if existing.is_file():
            existing.unlink()

    invoices, threshold, blocked = read_inputs(args.input_dir)
    if [invoice.invoice_id for invoice in invoices] != ["ACM-2026-0412", "GL-88210", "SP-0091", "AS-7744"]:
        raise ValueError("四张发票的人工审计顺序已变化")
    if threshold != Decimal("5000"):
        raise ValueError("有效复核阈值已变化")
    if sum((invoice.effective_amount for invoice in invoices if not invoice.blocked), Decimal("0")) != Decimal("16090.50"):
        raise ValueError("待处理发票合计与人工真值不一致")

    build_workbook(invoices, threshold, blocked, args.output_dir / XLSX_NAME)
    build_document(invoices, threshold, args.output_dir / DOCX_NAME)
    build_mbox(args.output_dir / MBOX_NAME)


if __name__ == "__main__":
    main()
