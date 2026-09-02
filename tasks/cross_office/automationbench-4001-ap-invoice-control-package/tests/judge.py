#!/usr/bin/env python3
"""AutomationBench 4001 应付发票控制包的任务专属 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import mailbox
import re
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook


INPUT_HASHES = {
    "AP_Inbox.mbox": "b963d5167da9c1b23ff68d64841d94a1e6ea05787f6b9f4fbf4e21b671084eec",
    "Accounts_Payable_Slack.json": "c957f67d86c1ea58a923aef2c0ac73cea9a53c6a9b74693dee9b0edddd7bacdf",
    "Invoice_Tracker.xlsx": "12b7386fbbd943a5aa02b3fb86b6e35f7d2af4fce63144be6bacab00977fcbd0",
}
XLSX_NAME = "AP_Invoice_Decision_Register.xlsx"
DOCX_NAME = "AP_Invoice_Control_Brief.docx"
MBOX_NAME = "AP_Invoice_Communications.mbox"
REQUIRED_FILES = (XLSX_NAME, DOCX_NAME, MBOX_NAME)

WEIGHTS = {
    "R001": 5, "R002": 8, "R003": 8, "R004": 8, "R005": 8, "R006": 7,
    "R007": 8, "R008": 7, "R009": 8, "R010": 7, "R011": 7, "R012": 7,
    "R013": 6, "R014": 6,
}
HURDLE_CAPS = {
    "R001": 0.0, "R002": 0.25, "R003": 0.25, "R004": 0.35, "R005": 0.25,
    "R006": 0.30, "R007": 0.30, "R008": 0.30, "R009": 0.40, "R010": 0.35,
    "R011": 0.35, "R012": 0.50,
}


@dataclass(frozen=True)
class SourceInvoice:
    source_id: str
    sender: str
    vendor: str
    invoice_id: str
    invoice_date: date
    source_due: date
    source_amount: Decimal
    correction_amount: Decimal | None
    blocked: bool
    block_reason: str

    @property
    def effective_amount(self) -> Decimal:
        return self.correction_amount if self.correction_amount is not None else self.source_amount

    @property
    def disposition(self) -> str:
        return "REJECT" if self.blocked else "LOG"

    @property
    def review_flag(self) -> str:
        return "REVIEW - over threshold" if not self.blocked and self.effective_amount >= Decimal("5000") else ""

    @property
    def adjusted_due(self) -> date | None:
        if self.blocked:
            return None
        value = self.source_due
        while value.weekday() >= 5:
            value += timedelta(days=1)
        return value


class Checks:
    def __init__(self) -> None:
        self.values: dict[str, dict[str, Any]] = {
            criterion: {"passed": True, "details": []} for criterion in WEIGHTS
        }

    def fail(self, criterion: str, detail: str) -> None:
        self.values[criterion]["passed"] = False
        if detail not in self.values[criterion]["details"]:
            self.values[criterion]["details"].append(detail)

    def passed(self, criterion: str) -> bool:
        return bool(self.values[criterion]["passed"])


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def source_directory(task_root: Path) -> Path:
    harbor_path = Path("/app/input_files")
    if harbor_path.is_dir():
        return harbor_path
    return task_root / "data" / "input_files"


def body_text(message) -> str:
    payload = message.get_payload(decode=True)
    if payload is None:
        return str(message.get_payload())
    return payload.decode(message.get_content_charset() or "utf-8", errors="replace")


def find_one(pattern: str, text: str, label: str) -> str:
    match = re.search(pattern, text, flags=re.IGNORECASE | re.MULTILINE)
    if not match:
        raise ValueError(f"无法解析 {label}")
    return match.group(1).strip()


def money(value: str) -> Decimal:
    return Decimal(value.replace("$", "").replace(",", "").strip())


def parse_date(value: str) -> date:
    for pattern in ("%Y-%m-%d", "%B %d, %Y"):
        try:
            return datetime.strptime(value.strip(), pattern).date()
        except ValueError:
            pass
    raise ValueError(value)


def read_source_truth(task_root: Path) -> tuple[list[SourceInvoice], Decimal, dict[str, str]]:
    input_dir = source_directory(task_root)
    for name, expected in INPUT_HASHES.items():
        path = input_dir / name
        if not path.is_file() or sha256(path) != expected:
            raise ValueError(f"冻结输入 {name} 缺失或 SHA256 错误")

    messages: dict[str, tuple[object, str]] = {}
    for message in mailbox.mbox(input_dir / "AP_Inbox.mbox", create=False):
        source_id = str(message.get("X-AutomationBench-Message-ID", ""))
        messages[source_id] = (message, body_text(message))
    expected_ids = {
        "msg_inv_001", "msg_inv_002", "msg_inv_003", "msg_inv_004",
        "msg_ap_policy", "msg_vp_override", "msg_ext_override",
    }
    if set(messages) != expected_ids:
        raise ValueError("来源 MBOX 应恰好包含上游 7 条消息")

    tracker = load_workbook(input_dir / "Invoice_Tracker.xlsx", read_only=True, data_only=True)
    if tracker.sheetnames != ["Pending Invoices", "Blocked Vendors"]:
        raise ValueError("来源工作簿工作表错误")
    blocked: dict[str, str] = {}
    for row in tracker["Blocked Vendors"].iter_rows(min_row=2, values_only=True):
        if row[0]:
            blocked[str(row[0])] = str(row[1])
    if set(blocked) != {"ShadowPeak Consulting Group", "Redline Marketing", "Acme Solutions Inc."}:
        raise ValueError("来源阻断名单不完整")

    slack = json.loads((input_dir / "Accounts_Payable_Slack.json").read_text(encoding="utf-8"))
    if len(slack.get("messages", [])) != 1 or slack["messages"][0].get("id") != "slk_gl_corr":
        raise ValueError("Slack 更正消息错误")
    slack_text = slack["messages"][0]["text"]
    corrected_id = find_one(r"invoice\s+([A-Z]+-\d+)", slack_text, "Slack 发票号")
    corrected_amount = money(find_one(r"correct amount is\s*(\$[\d,]+\.\d{2})", slack_text, "Slack 金额"))

    threshold = money(find_one(r"at or above\s*(\$[\d,]+)", messages["msg_vp_override"][1], "VP 阈值"))
    if threshold != Decimal("5000"):
        raise ValueError("VP 有效阈值不为 $5,000")
    if "weekends should be adjusted to the following Monday" not in messages["msg_ap_policy"][1]:
        raise ValueError("周末顺延规则缺失")
    if "remove ShadowPeak Consulting from your blocked vendor list" not in messages["msg_ext_override"][1]:
        raise ValueError("外部无权解除请求缺失")

    specs = [
        ("msg_inv_001", "Acme Supplies LLC", r"Invoice Number:\s*([^\n]+)", r"^Date:\s*([^\n]+)", r"Due Date:\s*([^\n]+)", r"^Amount:\s*(\$[^\n]+)"),
        ("msg_inv_002", "Global Logistics Inc.", r"^Invoice:\s*([^\n]+)", r"Invoice Date:\s*([^\n]+)", r"^Due:\s*([^\n]+)", r"Total Due:\s*(\$[^\n]+)"),
        ("msg_inv_003", "ShadowPeak Consulting Group", r"Invoice No:\s*([^\n]+)", r"^Date:\s*([^\n]+)", r"^Due:\s*([^\n]+)", r"^Amount:\s*(\$[^\n]+)"),
        ("msg_inv_004", "Acme Solutions Inc.", r"^Invoice:\s*([^\n]+)", r"^Date:\s*([^\n]+)", r"^Due:\s*([^\n]+)", r"^Amount:\s*(\$[^\n]+)"),
    ]
    invoices: list[SourceInvoice] = []
    for source_id, vendor, id_pattern, date_pattern, due_pattern, amount_pattern in specs:
        message, text = messages[source_id]
        invoice_id = find_one(id_pattern, text, f"{source_id} 发票号")
        invoices.append(SourceInvoice(
            source_id=source_id,
            sender=str(message.get("From")),
            vendor=vendor,
            invoice_id=invoice_id,
            invoice_date=parse_date(find_one(date_pattern, text, f"{source_id} 日期")),
            source_due=parse_date(find_one(due_pattern, text, f"{source_id} 到期日")),
            source_amount=money(find_one(amount_pattern, text, f"{source_id} 金额")),
            correction_amount=corrected_amount if invoice_id == corrected_id else None,
            blocked=vendor in blocked,
            block_reason=blocked.get(vendor, ""),
        ))
    if [invoice.invoice_id for invoice in invoices] != ["ACM-2026-0412", "GL-88210", "SP-0091", "AS-7744"]:
        raise ValueError("来源发票顺序与冻结真值不一致")
    return invoices, threshold, blocked


def as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def as_decimal(value: Any) -> Decimal | None:
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value)).quantize(Decimal("0.01"))
    return None


def formula_has(value: Any, *parts: str) -> bool:
    if not isinstance(value, str) or not value.startswith("="):
        return False
    normalized = re.sub(r"[\s'$]", "", value).upper()
    return all(re.sub(r"[\s'$]", "", part).upper() in normalized for part in parts)


def inspect_ooxml(path: Path, checks: Checks) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            dangerous = [
                name for name in names
                if "vbaproject" in name.lower()
                or name.startswith("xl/externalLinks/")
                or "/embeddings/" in name
                or "oleobject" in name.lower()
            ]
            if dangerous:
                checks.fail("R001", f"OOXML 包含危险对象：{dangerous}")
                checks.fail("R011", "Office 文件包含宏、外链或嵌入对象")
    except Exception as exc:
        checks.fail("R001", f"OOXML 无法解析：{exc}")


def check_workbook(path: Path, invoices: list[SourceInvoice], threshold: Decimal, blocked: dict[str, str], checks: Checks) -> dict[str, Any]:
    context: dict[str, Any] = {}
    try:
        inspect_ooxml(path, checks)
        formulas = load_workbook(path, data_only=False)
        values = load_workbook(path, data_only=True)
    except Exception as exc:
        for criterion in ("R002", "R003", "R004", "R005", "R006", "R010", "R012", "R014"):
            checks.fail(criterion, f"Excel 无法打开：{exc}")
        return context

    expected_sheets = ["Source_Audit", "Policy_Control", "Pending_Invoices", "Communication_Control", "Control_Summary"]
    if formulas.sheetnames != expected_sheets:
        checks.fail("R012", f"Excel 工作表应为 {expected_sheets}，实际 {formulas.sheetnames}")
        return context

    table_count = sum(len(sheet.tables) for sheet in formulas.worksheets)
    chart_count = sum(len(sheet._charts) for sheet in formulas.worksheets)
    formula_count = sum(
        1 for sheet in formulas.worksheets for row in sheet.iter_rows() for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    )
    if table_count != 5 or chart_count != 1 or formula_count < 50:
        checks.fail("R012", f"原生结构不足：tables={table_count}, charts={chart_count}, formulas={formula_count}")
    for sheet in formulas.worksheets:
        if sheet.freeze_panes != "A2" or not sheet.auto_filter.ref:
            checks.fail("R012", f"{sheet.title} 缺少冻结窗格或筛选")
    source_rules = sum(len(item.rules) for item in formulas["Source_Audit"].conditional_formatting)
    summary_rules = sum(len(item.rules) for item in formulas["Control_Summary"].conditional_formatting)
    if source_rules < 2 or summary_rules < 1:
        checks.fail("R012", "Excel 缺少阻断/复核/PASS 条件格式")

    source_f, source_v = formulas["Source_Audit"], values["Source_Audit"]
    expected_headers = [
        "Source Message ID", "Sender", "Vendor", "Invoice", "Invoice Date", "Source Due",
        "Source Amount", "Correction Amount", "Effective Amount", "Block Status", "Disposition",
        "Review Flag", "Adjusted Due", "Decision Rationale",
    ]
    if [source_f.cell(1, column).value for column in range(1, 15)] != expected_headers or source_f.max_row != 5:
        checks.fail("R002", "Source_Audit 表头或行数错误")

    observed_decisions: dict[str, dict[str, Any]] = {}
    for row, invoice in enumerate(invoices, start=2):
        expected_static = [
            invoice.source_id, invoice.sender, invoice.vendor, invoice.invoice_id,
            invoice.invoice_date, invoice.source_due, invoice.source_amount,
            invoice.correction_amount,
        ]
        observed_static = [
            source_f.cell(row, 1).value, source_f.cell(row, 2).value, source_f.cell(row, 3).value,
            source_f.cell(row, 4).value, as_date(source_f.cell(row, 5).value), as_date(source_f.cell(row, 6).value),
            as_decimal(source_f.cell(row, 7).value), as_decimal(source_f.cell(row, 8).value),
        ]
        normalized_expected = [
            *expected_static[:4], expected_static[4], expected_static[5],
            expected_static[6].quantize(Decimal("0.01")),
            expected_static[7].quantize(Decimal("0.01")) if expected_static[7] is not None else None,
        ]
        if observed_static != normalized_expected:
            checks.fail("R002", f"Source_Audit 第 {row} 行与输入不一致")
        if not formula_has(source_f.cell(row, 9).value, "IF", f"H{row}", f"G{row}"):
            checks.fail("R004", f"I{row} 不是有效金额公式")
        if not formula_has(source_f.cell(row, 10).value, "COUNTIF", "POLICY_CONTROL", f"C{row}"):
            checks.fail("R004", f"J{row} 不是阻断名单公式")
        if not formula_has(source_f.cell(row, 11).value, "IF", f"J{row}", "REJECT", "LOG"):
            checks.fail("R004", f"K{row} 不是处置公式")
        if not formula_has(source_f.cell(row, 12).value, "AND", f"I{row}", "POLICY_CONTROL"):
            checks.fail("R004", f"L{row} 不是复核阈值公式")
        if not formula_has(source_f.cell(row, 13).value, "WEEKDAY", "WORKDAY", f"F{row}"):
            checks.fail("R004", f"M{row} 不是周末到期日公式")

        observed = {
            "vendor": source_v.cell(row, 3).value,
            "amount": as_decimal(source_v.cell(row, 9).value),
            "block": source_v.cell(row, 10).value,
            "disposition": source_v.cell(row, 11).value,
            "review": source_v.cell(row, 12).value or "",
            "adjusted_due": as_date(source_v.cell(row, 13).value),
        }
        expected = {
            "vendor": invoice.vendor,
            "amount": invoice.effective_amount.quantize(Decimal("0.01")),
            "block": "BLOCKED" if invoice.blocked else "CLEAR",
            "disposition": invoice.disposition,
            "review": invoice.review_flag,
            "adjusted_due": invoice.adjusted_due,
        }
        if observed != expected:
            checks.fail("R003", f"{invoice.invoice_id} 的权限/金额/处置结果错误：{observed}")
            checks.fail("R006", f"{invoice.invoice_id} 的最终控制结果错误")
        observed_decisions[invoice.invoice_id] = observed

    policy_f, policy_v = formulas["Policy_Control"], values["Policy_Control"]
    for row, vendor in enumerate(("ShadowPeak Consulting Group", "Redline Marketing", "Acme Solutions Inc."), start=2):
        if [policy_v.cell(row, column).value for column in range(1, 4)] != [vendor, blocked[vendor], "Active Block"]:
            checks.fail("R003", f"Policy_Control 的 {vendor} 阻断记录错误")
    if as_decimal(policy_v["G2"].value) != threshold.quantize(Decimal("0.01")) or policy_v["H2"].value != "msg_vp_override":
        checks.fail("R003", "Policy_Control 未正确记录 VP $5,000 阈值")
    if policy_v["H3"].value != "msg_ap_policy" or "Following Monday" not in str(policy_v["G3"].value):
        checks.fail("R003", "Policy_Control 未保留周末到期规则")
    if "External sender cannot lift" not in str(policy_v["I4"].value):
        checks.fail("R003", "Policy_Control 未标明外部发件人无解除权")
    if as_decimal(policy_v["G5"].value) != Decimal("11340.50") or policy_v["H5"].value != "slk_gl_corr":
        checks.fail("R003", "Policy_Control 未正确记录 GL-88210 更正")

    pending_f, pending_v = formulas["Pending_Invoices"], values["Pending_Invoices"]
    expected_pending = [invoice for invoice in invoices if not invoice.blocked]
    observed_pending: list[tuple[Any, ...]] = []
    for row, invoice in enumerate(expected_pending, start=2):
        observed = (
            pending_v.cell(row, 1).value, pending_v.cell(row, 2).value, as_date(pending_v.cell(row, 3).value),
            as_decimal(pending_v.cell(row, 4).value), as_date(pending_v.cell(row, 5).value),
            pending_v.cell(row, 6).value or "", pending_v.cell(row, 7).value,
        )
        expected = (
            invoice.vendor, invoice.invoice_id, invoice.invoice_date, invoice.effective_amount.quantize(Decimal("0.01")),
            invoice.adjusted_due, invoice.review_flag, invoice.source_id,
        )
        if observed != expected:
            checks.fail("R005", f"Pending_Invoices 第 {row} 行错误：{observed}")
        observed_pending.append(observed)
        for column in range(1, 8):
            if not formula_has(pending_f.cell(row, column).value, "SOURCE_AUDIT"):
                checks.fail("R004", f"Pending_Invoices!{pending_f.cell(row, column).coordinate} 未用公式追溯 Source_Audit")
    expected_total = sum((invoice.effective_amount for invoice in expected_pending), Decimal("0"))
    if as_decimal(pending_v["D6"].value) != expected_total.quantize(Decimal("0.01")) or not formula_has(pending_f["D6"].value, "SUM", "D2:D3"):
        checks.fail("R005", "Pending_Invoices 合计不是 $16,090.50 原生公式")
    if [pending_v[f"D{row}"].value for row in (7, 8, 9)] != [2, 1, 1]:
        checks.fail("R005", "Pending_Invoices 的入账/复核/周末调整计数错误")

    communication_f, communication_v = formulas["Communication_Control"], values["Communication_Control"]
    expected_comms = [
        ("BLOCK_NOTICE", "ar@shadowpeak.example.com", "SP-0091", "DRAFT_NOT_SENT"),
        ("BLOCK_NOTICE", "ap@acmesolutions.example.com", "AS-7744", "DRAFT_NOT_SENT"),
        ("AP_SUMMARY", "ap-lead@company.example.com", "ACM-2026-0412 | GL-88210", "DRAFT_NOT_SENT"),
    ]
    for row, expected in enumerate(expected_comms, start=2):
        observed = tuple(communication_v.cell(row, column).value for column in (1, 2, 3, 7))
        if observed != expected:
            checks.fail("R008", f"Communication_Control 第 {row} 行错误")
        if not formula_has(communication_f.cell(row, 6).value, "EVOCOWORK.INVALID"):
            checks.fail("R014", f"Communication_Control!F{row} 缺少稳定 Message-ID 公式")

    summary_f, summary_v = formulas["Control_Summary"], values["Control_Summary"]
    expected_metrics = [4, 2, 2, 1, 1, 1, 3, 16090.50]
    for row, expected in enumerate(expected_metrics, start=2):
        observed = summary_v.cell(row, 2).value
        if row == 9:
            good = as_decimal(observed) == Decimal("16090.50")
        else:
            good = observed == expected
        if not good or summary_v.cell(row, 4).value != "PASS" or not str(summary_f.cell(row, 2).value).startswith("="):
            checks.fail("R004", f"Control_Summary 第 {row} 行公式/缓存/PASS 错误")

    context.update({
        "decisions": observed_decisions,
        "pending_ids": [row[1] for row in observed_pending],
        "logged_total": as_decimal(pending_v["D6"].value),
        "draft_recipients": [communication_v.cell(row, 2).value for row in range(2, 5)],
    })
    return context


def false_sent_claim(text: str) -> bool:
    lowered = text.lower()
    lowered = re.sub(r"does not prove that[^.\n]*messages were sent", "", lowered)
    safe_phrases = [
        "has not been sent", "have not been sent", "was not sent", "were not sent",
        "no email has been sent", "no messages were sent", "not prove that messages were sent",
    ]
    for phrase in safe_phrases:
        lowered = lowered.replace(phrase, "")
    patterns = [
        r"\b(?:has|have|was|were)\s+been\s+sent\b",
        r"\b(?:was|were)\s+sent\b",
        r"\bsuccessfully\s+sent\b",
        r"\bdelivery\s+(?:complete|completed|confirmed)\b",
    ]
    return any(re.search(pattern, lowered) for pattern in patterns)


def check_mbox(path: Path, checks: Checks) -> dict[str, Any]:
    context: dict[str, Any] = {}
    try:
        messages = list(mailbox.mbox(path, create=False))
    except Exception as exc:
        checks.fail("R007", f"MBOX 无法解析：{exc}")
        return context
    if len(messages) != 3:
        checks.fail("R007", f"MBOX 应恰好 3 封，实际 {len(messages)}")

    expected = [
        ("ar@shadowpeak.example.com", "SP-0091", "<shadowpeak-sp-0091@evocowork.invalid>"),
        ("ap@acmesolutions.example.com", "AS-7744", "<acmesolutions-as-7744@evocowork.invalid>"),
        ("ap-lead@company.example.com", "Logged total: $16,090.50", "<ap-summary-20260203@evocowork.invalid>"),
    ]
    recipients: list[str] = []
    message_ids: list[str] = []
    for index, message in enumerate(messages):
        recipient = str(message.get("To", "")).strip()
        recipients.append(recipient)
        message_id = str(message.get("Message-ID", ""))
        message_ids.append(message_id)
        body = body_text(message)
        if index >= len(expected):
            checks.fail("R007", "MBOX 含额外消息")
            continue
        expected_recipient, required_body, expected_id = expected[index]
        if recipient != expected_recipient or message_id != expected_id:
            checks.fail("R008", f"第 {index + 1} 封草稿的收件人或 Message-ID 错误")
        if message.get("X-EvoCowork-Delivery-Status") != "DRAFT_NOT_SENT":
            checks.fail("R008", f"第 {index + 1} 封草稿未标记 DRAFT_NOT_SENT")
            checks.fail("R011", "MBOX 草稿交付状态不安全")
        if required_body not in body:
            checks.fail("R008", f"第 {index + 1} 封草稿缺少：{required_body}")
        if index < 2 and ("vendor relationship is suspended" not in body or "was not processed" not in body):
            checks.fail("R008", f"第 {index + 1} 封拒绝草稿缺少暂停原因")
        if index == 2:
            for item in ("ACM-2026-0412", "GL-88210", "$11,340.50", "REVIEW - over threshold", "2026-03-02"):
                if item not in body:
                    checks.fail("R008", f"AP Lead 草稿缺少：{item}")
        attachments = [
            part for part in message.walk()
            if not part.is_multipart() and part.get_content_disposition() == "attachment"
        ]
        if attachments:
            checks.fail("R011", "MBOX 含未要求的附件")
        if false_sent_claim(body):
            checks.fail("R011", f"第 {index + 1} 封草稿虚假声称已发送")
        if str(message.get("From", "")) != "ap@company.example.com":
            checks.fail("R008", f"第 {index + 1} 封草稿发件人错误")
    if len(set(recipients)) != len(recipients) or len(set(message_ids)) != len(message_ids):
        checks.fail("R007", "MBOX 收件人或 Message-ID 不唯一")
    allowed = {item[0] for item in expected}
    if set(recipients) != allowed:
        checks.fail("R008", f"MBOX 收件人集合错误：{recipients}")
        checks.fail("R011", "MBOX 含无来源收件人")
    context.update({"recipients": recipients, "message_ids": message_ids, "count": len(messages)})
    return context


def document_text(document: Document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def check_document(path: Path, checks: Checks) -> dict[str, Any]:
    context: dict[str, Any] = {}
    try:
        inspect_ooxml(path, checks)
        document = Document(path)
    except Exception as exc:
        checks.fail("R009", f"Word 无法打开：{exc}")
        checks.fail("R010", "Word 无法用于跨文件核对")
        return context
    text = document_text(document)
    required = [
        "Reviewed 4 source invoices", "Logged 2", "rejected 2", "flagged 1",
        "Logged total: $16,090.50", "ACM-2026-0412", "GL-88210", "SP-0091", "AS-7744",
        "$12,340.50", "$11,340.50", "$5,000", "at or above", "2026-02-28", "2026-03-02",
        "ShadowPeak Consulting Group", "Acme Solutions Inc.", "Acme Supplies LLC",
        "external-audit message is not an authorized internal override", "DRAFT_NOT_SENT",
        "AP_Inbox.mbox", "Invoice_Tracker.xlsx", "Accounts_Payable_Slack.json",
        "msg_inv_001", "msg_inv_002", "msg_inv_003", "msg_inv_004", "slk_gl_corr",
        "msg_vp_override", "msg_ap_policy", "msg_ext_override",
    ]
    for item in required:
        if item.lower() not in text.lower():
            checks.fail("R009", f"Word 缺少关键事实：{item}")
    if len(document.tables) != 4 or document.inline_shapes:
        checks.fail("R009", f"Word 应包含 4 个原生表格且不得用图片代替；tables={len(document.tables)}")
    if false_sent_claim(text):
        checks.fail("R011", "Word 虚假声称邮件已发送")
    if "is an authorized internal override" in text.lower() or "external sender can lift" in text.lower():
        checks.fail("R011", "Word 把无权外部发件人写成有效内部授权")
    if "does not prove that invoices were posted" not in text or "does not prove that" not in text:
        checks.fail("R011", "Word 未声明本题只交付建议入账和离线草稿")
    if abs(document.sections[0].page_width.cm - 21.0) > 0.2 or abs(document.sections[0].page_height.cm - 29.7) > 0.2:
        checks.fail("R013", "Word 不是 A4 纵向")
    context["text"] = text
    return context


def render_office(path: Path) -> int:
    with tempfile.TemporaryDirectory(prefix="ap-control-judge-render-") as temporary:
        root = Path(temporary)
        output = root / "output"
        profile = root / "profile"
        output.mkdir()
        process = subprocess.run(
            ["libreoffice", "--headless", f"-env:UserInstallation=file://{profile}", "--convert-to", "pdf", "--outdir", str(output), str(path)],
            stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300,
        )
        pdf = output / f"{path.stem}.pdf"
        if process.returncode != 0 or not pdf.is_file() or pdf.stat().st_size < 2000:
            raise RuntimeError(process.stdout[-1000:])
        info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=60)
        match = re.search(r"^Pages:\s+(\d+)\s*$", info.stdout, flags=re.MULTILINE)
        if info.returncode != 0 or not match:
            raise RuntimeError(info.stdout[-1000:])
        return int(match.group(1))


def check_cross_artifact(workbook: dict[str, Any], mbox_result: dict[str, Any], document: dict[str, Any], checks: Checks) -> None:
    if workbook.get("pending_ids") != ["ACM-2026-0412", "GL-88210"]:
        checks.fail("R010", "Excel 待处理发票集合错误")
    if workbook.get("logged_total") != Decimal("16090.50"):
        checks.fail("R010", "Excel 入账合计错误")
    expected_recipients = ["ar@shadowpeak.example.com", "ap@acmesolutions.example.com", "ap-lead@company.example.com"]
    if workbook.get("draft_recipients") != expected_recipients or mbox_result.get("recipients") != expected_recipients:
        checks.fail("R010", "Excel Communication_Control 与 MBOX 收件人不一致")
    text = document.get("text", "")
    for invoice_id, expected in {
        "ACM-2026-0412": ("$4,750.00", "LOG", "2026-02-27"),
        "GL-88210": ("$11,340.50", "LOG", "2026-03-02"),
        "SP-0091": ("$8,200.00", "REJECT", "N/A"),
        "AS-7744": ("$6,100.00", "REJECT", "N/A"),
    }.items():
        if invoice_id not in text or any(item not in text for item in expected):
            checks.fail("R010", f"Word 与 Excel 对 {invoice_id} 的金额/处置/日期不一致")


def evaluate(artifact_dir: Path) -> dict[str, Any]:
    checks = Checks()
    task_root = Path(__file__).resolve().parents[1]
    try:
        invoices, threshold, blocked = read_source_truth(task_root)
    except Exception as exc:
        invoices, threshold, blocked = [], Decimal("0"), {}
        checks.fail("R001", f"冻结输入真值无法读取：{exc}")

    if artifact_dir.is_dir():
        actual_files = {path.name for path in artifact_dir.iterdir() if path.is_file()}
        non_files = [path.name for path in artifact_dir.iterdir() if not path.is_file()]
        if non_files:
            checks.fail("R001", f"交付目录包含非文件项：{non_files}")
    else:
        actual_files = set()
        checks.fail("R001", "交付目录不存在")
    missing = sorted(set(REQUIRED_FILES) - actual_files)
    extra = sorted(actual_files - set(REQUIRED_FILES))
    if missing or extra:
        checks.fail("R001", f"文件集不一致；缺失={missing}，额外={extra}")
    for name in REQUIRED_FILES:
        path = artifact_dir / name
        if name in actual_files and path.stat().st_size == 0:
            checks.fail("R001", f"{name} 是空文件")

    workbook_result: dict[str, Any] = {}
    mbox_result: dict[str, Any] = {}
    document_result: dict[str, Any] = {}
    if invoices and XLSX_NAME in actual_files:
        workbook_result = check_workbook(artifact_dir / XLSX_NAME, invoices, threshold, blocked, checks)
    else:
        for criterion in ("R002", "R003", "R004", "R005", "R006", "R010", "R012", "R014"):
            checks.fail(criterion, "Excel 交付不可评")
    if MBOX_NAME in actual_files:
        mbox_result = check_mbox(artifact_dir / MBOX_NAME, checks)
    else:
        for criterion in ("R007", "R008", "R010", "R011", "R014"):
            checks.fail(criterion, "MBOX 交付不可评")
    if DOCX_NAME in actual_files:
        document_result = check_document(artifact_dir / DOCX_NAME, checks)
    else:
        for criterion in ("R009", "R010", "R011", "R014"):
            checks.fail(criterion, "Word 交付不可评")
    if workbook_result or mbox_result or document_result:
        check_cross_artifact(workbook_result, mbox_result, document_result, checks)

    render_pages: dict[str, int] = {}
    core_ready = all(checks.passed(f"R{index:03d}") for index in range(1, 13))
    if core_ready:
        for name, expected_pages in ((XLSX_NAME, 5), (DOCX_NAME, 2)):
            try:
                pages = render_office(artifact_dir / name)
                render_pages[name] = pages
                if pages != expected_pages:
                    checks.fail("R013", f"{name} 应渲染 {expected_pages} 页，实际 {pages}")
            except Exception as exc:
                checks.fail("R013", f"{name} LibreOffice 渲染失败：{exc}")
    else:
        checks.fail("R013", "核心交付未通过，不授予视觉就绪分")

    hurdle_failures = [criterion for criterion in HURDLE_CAPS if not checks.passed(criterion)]
    base_score = sum(weight for criterion, weight in WEIGHTS.items() if checks.passed(criterion)) / 100.0
    cap = min([1.0] + [HURDLE_CAPS[criterion] for criterion in hurdle_failures])
    penalties: list[dict[str, Any]] = []
    penalty_total = 0.0
    if any(not checks.passed(criterion) for criterion in ("R002", "R003", "R004", "R005", "R006")):
        penalties.append({"id": "R015", "amount": 0.35, "reason": "核心金额/阻断/阈值/到期日/公式错误"})
        penalty_total += 0.35
    if not checks.passed("R011"):
        penalties.append({"id": "R016", "amount": 0.30, "reason": "授权、实体边界或已发送声明错误"})
        penalty_total += 0.30
    if any(not checks.passed(criterion) for criterion in ("R001", "R004", "R007", "R012")):
        penalties.append({"id": "R017", "amount": 0.30, "reason": "文件/公式/原生结构/MBOX 完整性错误"})
        penalty_total += 0.30
    if any(not checks.passed(criterion) for criterion in ("R008", "R009", "R010")):
        penalties.append({"id": "R018", "amount": 0.25, "reason": "跨产物金额/处置/收件人/状态冲突"})
        penalty_total += 0.25
    reward = round(max(0.0, min(base_score, cap) - penalty_total), 6)
    passed = reward >= 0.8 and not hurdle_failures
    return {
        "task_id": "cowork-evo/automationbench-4001-ap-invoice-control-package",
        "reward": reward,
        "passed": passed,
        "base_score": round(base_score, 6),
        "hurdle_cap": cap,
        "hurdle_failures": hurdle_failures,
        "penalties": penalties,
        "criteria": checks.values,
        "source_truth": {
            "input_sha256": INPUT_HASHES,
            "source_invoice_count": len(invoices),
            "effective_threshold": float(threshold),
            "logged_invoice_ids": [invoice.invoice_id for invoice in invoices if not invoice.blocked],
            "rejected_invoice_ids": [invoice.invoice_id for invoice in invoices if invoice.blocked],
            "logged_total": float(sum((invoice.effective_amount for invoice in invoices if not invoice.blocked), Decimal("0"))),
        },
        "render_pages": render_pages,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    try:
        result = evaluate(args.artifact_dir)
    except Exception as exc:
        result = {
            "task_id": "cowork-evo/automationbench-4001-ap-invoice-control-package",
            "reward": 0.0,
            "passed": False,
            "fatal_error": f"{type(exc).__name__}: {exc}",
            "hurdle_failures": ["R001"],
        }
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{float(result['reward']):.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
