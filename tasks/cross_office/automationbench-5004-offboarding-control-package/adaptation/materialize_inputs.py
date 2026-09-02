#!/usr/bin/env python3
"""只把 AutomationBench 5004 的公开初始状态静态化为本题输入。

本脚本只接受 ``hr.offboarding_automation``，并逐项断言四名员工、
四条政策、运行时间、Slack 频道和十三条上游断言。它不是通用制题器。
"""

from __future__ import annotations

import importlib.util
import json
import shutil
import sys
import types
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
UPSTREAM = Path(__file__).with_name("upstream-hr-tasks.py")
SNAPSHOT = Path(__file__).with_name("upstream-task-5004.json")
DATA_DIR = ROOT / "data" / "input_files"
ENV_DIR = ROOT / "environment" / "input_files"
FIXED_ZIP_TIME = (2026, 3, 10, 9, 0, 0)


def load_task() -> dict:
    datasets = types.ModuleType("datasets")
    datasets.Dataset = object
    sys.modules["datasets"] = datasets
    for name in ("automationbench", "automationbench.domains", "automationbench.domains.hr"):
        sys.modules[name] = types.ModuleType(name)
    noise = types.ModuleType("automationbench.domains.hr._noise")
    noise.apply_noise = lambda value: value
    sys.modules[noise.__name__] = noise

    spec = importlib.util.spec_from_file_location("automationbench_hr_tasks_frozen", UPSTREAM)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入冻结的 AutomationBench HR 源文件")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    task = module.get_hr_offboarding_task()
    if task.get("example_id") != 5004 or task.get("task") != "hr.offboarding_automation":
        raise ValueError("冻结函数不再是 AutomationBench 5004")
    return task


def assert_shape(task: dict) -> None:
    state = task["info"]["initial_state"]
    assert state["meta"]["current_time"] == "2026-03-10T09:00:00Z"
    books = state["google_sheets"]["spreadsheets"]
    assert len(books) == 1 and books[0]["id"] == "ss_offboard"
    sheets = books[0]["worksheets"]
    assert [sheet["id"] for sheet in sheets] == ["ws_departures", "ws_policy"]
    departures = sheets[0]["rows"]
    assert [row["cells"]["Employee"] for row in departures] == [
        "Greg Foster", "Diana Reese", "Raj Patel", "Greg Forster"
    ]
    assert [row["row_id"] for row in departures] == [2, 3, 4, 5]
    policies = sheets[1]["rows"]
    assert [row["cells"]["Policy"] for row in policies] == [
        "Voluntary Resignation",
        "Involuntary Termination",
        "Future-Dated Departures",
        "Severance Payments",
    ]
    assert state["slack"]["channels"] == [
        {"id": "C_DEPART", "name": "departures", "is_private": False}
    ]
    assert len(task["info"]["assertions"]) == 13


def normalize_xlsx(path: Path) -> None:
    temporary = path.with_suffix(".normalized.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        temporary, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as destination:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            destination.writestr(info, source.read(name))
    temporary.replace(path)


def format_sheet(sheet, widths: dict[str, int]) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = sheet.dimensions


def write_workbook(state: dict, target: Path) -> None:
    source = state["google_sheets"]["spreadsheets"][0]
    departures_source, policy_source = source["worksheets"]
    workbook = Workbook()
    workbook.creator = "AutomationBench 5004 public fixture"
    workbook.title = source["title"]
    workbook.subject = "Static Office representation of the offboarding source workbook"
    workbook.description = "No employee, date, status, note or policy was added beyond AutomationBench 5004."

    departures = workbook.active
    departures.title = departures_source["title"]
    departure_headers = [
        "Source Row", "Employee", "Employee Email", "Department", "Manager",
        "Manager Email", "Last Day", "Type", "Status", "Notes",
    ]
    departures.append(departure_headers)
    for item in departures_source["rows"]:
        cells = item["cells"]
        departures.append([item["row_id"]] + [cells.get(header, "") for header in departure_headers[1:]])
    table = Table(displayName="DepartureSourceInput", ref=departures.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    departures.add_table(table)
    format_sheet(departures, {
        "A": 12, "B": 22, "C": 34, "D": 18, "E": 22,
        "F": 34, "G": 16, "H": 34, "I": 16, "J": 42,
    })

    policy = workbook.create_sheet(policy_source["title"])
    policy.append(["Source Row", "Policy", "Steps"])
    for item in policy_source["rows"]:
        policy.append([item["row_id"], item["cells"]["Policy"], item["cells"]["Steps"]])
    policy_table = Table(displayName="OffboardingPolicyInput", ref=policy.dimensions)
    policy_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    policy.add_table(policy_table)
    format_sheet(policy, {"A": 12, "B": 30, "C": 110})
    for row in range(2, policy.max_row + 1):
        policy.row_dimensions[row].height = 54

    workbook.save(target)
    normalize_xlsx(target)


def main() -> None:
    task = load_task()
    assert_shape(task)
    state = task["info"]["initial_state"]
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    ENV_DIR.mkdir(parents=True, exist_ok=True)

    SNAPSHOT.write_text(json.dumps(task, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_workbook(state, DATA_DIR / "Employee_Departures.xlsx")
    (DATA_DIR / "Workflow_Context.json").write_text(
        json.dumps(state["meta"], ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (DATA_DIR / "Departures_Channel.json").write_text(
        json.dumps(state["slack"], ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    for source in sorted(DATA_DIR.iterdir()):
        if source.is_file():
            shutil.copy2(source, ENV_DIR / source.name)


if __name__ == "__main__":
    main()
