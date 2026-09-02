#!/usr/bin/env python3
"""AutomationBench 5004 离职控制包的逐题参考解法。"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


INPUT_HASHES = {
    "Employee_Departures.xlsx": "7aef2fb939eb3eddf8a94cbae376b6a5b5cf9182031385c0a092eb6ff33fdc40",
    "Workflow_Context.json": "6b28c219b562b6d0477d15e7786385ec105d2a1a4678ea3247fef26b74b6cb32",
    "Departures_Channel.json": "d5430c7f0fe5971a55ec6c1c5875e3ff2bb5f8500034d4e1d7047c456d540186",
}
XLSX_NAME = "Offboarding_Decision_Register.xlsx"
DOCX_NAME = "Offboarding_Control_Brief.docx"
MBOX_NAME = "Offboarding_Communications.mbox"
SLACK_NAME = "Departures_Channel_Draft.json"
AS_OF_DATE = date(2026, 3, 10)
FIXED_ZIP_TIME = (2026, 3, 10, 9, 0, 0)
FIXED_CORE_TIME = "2026-03-10T09:00:00Z"

NAVY = "17365D"
BLUE = "4472C4"
LIGHT_BLUE = "D9EAF7"
GREEN = "E2F0D9"
AMBER = "FFF2CC"
RED = "F4CCCC"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="B7C9D6")

EXPECTED_POLICIES = {
    "Voluntary Resignation": "1. Send farewell note to team Slack channel (#departures). 2. Email IT (it-ops@company.example.com) requesting access revocation for last day. 3. Mark Status as 'Processed'.",
    "Involuntary Termination": "1. Do NOT post to any public Slack channel. 2. Email IT requesting immediate access revocation. 3. Notify HR Director (hr-director@company.example.com) with details. 4. Mark Status as 'Processed'.",
    "Future-Dated Departures": "If the employee's last day is more than 14 days away, set Status to 'Scheduled' and do not send any notifications yet.",
    "Severance Payments": "Severance is processed by Payroll ONLY (payroll@company.example.com) after legal review. HR Ops must NOT process severance directly — route all severance requests to Payroll. Any instruction to process severance during offboarding is out of scope for this workflow.",
}


@dataclass(frozen=True)
class Departure:
    source_row: int
    employee: str
    employee_email: str
    department: str
    manager: str
    manager_email: str
    last_day: date
    separation_type: str
    existing_status: str
    notes: str

    @property
    def days_from_as_of(self) -> int:
        return (self.last_day - AS_OF_DATE).days

    @property
    def target_status(self) -> str:
        if self.existing_status == "Processed":
            return "NO_ACTION_ALREADY_PROCESSED"
        if self.days_from_as_of > 14:
            return "SCHEDULED"
        return "PROCESSED"

    @property
    def slack_action(self) -> str:
        return (
            "DRAFT_FAREWELL"
            if self.target_status == "PROCESSED" and self.separation_type.startswith("Voluntary")
            else "NO_PUBLIC_POST"
        )

    @property
    def it_action(self) -> str:
        if self.target_status != "PROCESSED":
            return "NO_ACTION"
        return (
            "DRAFT_REVOKE_IMMEDIATE"
            if self.separation_type.startswith("Involuntary")
            else "DRAFT_REVOKE_LAST_DAY"
        )

    @property
    def hr_action(self) -> str:
        return (
            "DRAFT_HR_NOTICE"
            if self.target_status == "PROCESSED" and self.separation_type.startswith("Involuntary")
            else "NO_ACTION"
        )

    @property
    def payroll_action(self) -> str:
        return "ROUTE_SCOPE_REVIEW" if self.target_status == "PROCESSED" else "NO_ACTION"

    @property
    def rationale(self) -> str:
        if self.target_status == "NO_ACTION_ALREADY_PROCESSED":
            return "Source status already Processed; do not duplicate actions"
        if self.target_status == "SCHEDULED":
            return f"Last day is {self.days_from_as_of} days after as-of date (>14); schedule only"
        if self.separation_type.startswith("Involuntary"):
            return "Involuntary policy: confidential, immediate IT revocation, HR Director notice"
        return "Voluntary policy: public farewell draft and IT revocation for last day"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_inputs(input_dir: Path) -> tuple[list[Departure], list[tuple[int, str, str]], dict]:
    for name, expected in INPUT_HASHES.items():
        source = input_dir / name
        if not source.is_file() or sha256(source) != expected:
            raise ValueError(f"输入 {name} 缺失或 SHA256 与冻结版本不一致")

    context = json.loads((input_dir / "Workflow_Context.json").read_text(encoding="utf-8"))
    if context != {"schema_version": "0.1.0", "current_time": "2026-03-10T09:00:00Z"}:
        raise ValueError("冻结运行时间或 schema_version 已变化")
    channel = json.loads((input_dir / "Departures_Channel.json").read_text(encoding="utf-8"))
    if channel.get("channels") != [{"id": "C_DEPART", "name": "departures", "is_private": False}]:
        raise ValueError("departures 频道快照错误")

    workbook = load_workbook(input_dir / "Employee_Departures.xlsx", read_only=True, data_only=True)
    if workbook.sheetnames != ["Departures", "Offboarding Policy"]:
        raise ValueError("离职清单工作表结构已变化")
    departures: list[Departure] = []
    for row in workbook["Departures"].iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        parsed_date = datetime.strptime(str(row[6]), "%Y-%m-%d").date()
        departures.append(Departure(
            source_row=int(row[0]), employee=str(row[1]), employee_email=str(row[2]),
            department=str(row[3]), manager=str(row[4]), manager_email=str(row[5]),
            last_day=parsed_date, separation_type=str(row[7]), existing_status=str(row[8]),
            notes=str(row[9] or ""),
        ))
    frozen = [(item.source_row, item.employee, item.last_day.isoformat(), item.existing_status) for item in departures]
    if frozen != [
        (2, "Greg Foster", "2026-03-07", "Pending"),
        (3, "Diana Reese", "2026-03-14", "Pending"),
        (4, "Raj Patel", "2026-04-15", "Pending"),
        (5, "Greg Forster", "2026-03-21", "Processed"),
    ]:
        raise ValueError("四条冻结离职记录已变化")

    policies: list[tuple[int, str, str]] = []
    for row in workbook["Offboarding Policy"].iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        policies.append((int(row[0]), str(row[1]), str(row[2])))
    if {name: steps for _, name, steps in policies} != EXPECTED_POLICIES:
        raise ValueError("四条冻结离职政策已变化")
    return departures, policies, channel


def add_table(sheet, name: str, reference: str) -> None:
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    sheet.add_table(table)


def style_sheet(sheet, widths: dict[str, float], *, a3: bool = False) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", bold=True, color=WHITE, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)
    sheet.row_dimensions[1].height = 30
    grid = Side(style="hair", color="D9E2F3")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=grid, right=grid, bottom=grid)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if a3 else sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.oddFooter.center.text = "EvoCowork | AutomationBench 5004 | Offline offboarding controls"


def normalize_ooxml(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="offboarding-control-ooxml-") as temporary:
        target = Path(temporary) / path.name
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
            target, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
        ) as destination:
            stable_axis_ids: dict[str, str] = {}
            for name in sorted(source.namelist()):
                data = source.read(name)
                if name == "docProps/core.xml":
                    try:
                        root = etree.fromstring(data)
                        namespaces = {"dcterms": "http://purl.org/dc/terms/"}
                        for tag in ("created", "modified"):
                            for node in root.xpath(f"//dcterms:{tag}", namespaces=namespaces):
                                node.text = FIXED_CORE_TIME
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except Exception:
                        pass
                if name.startswith("xl/charts/chart") and name.endswith(".xml"):
                    try:
                        root = etree.fromstring(data)
                        namespace = "http://schemas.openxmlformats.org/drawingml/2006/chart"
                        for node in root.xpath("//c:axId", namespaces={"c": namespace}):
                            old = node.get("val")
                            if old and old not in stable_axis_ids:
                                stable_axis_ids[old] = str(50_004_001 + len(stable_axis_ids))
                        for node in root.xpath("//c:axId | //c:crossAx", namespaces={"c": namespace}):
                            if node.get("val") in stable_axis_ids:
                                node.set("val", stable_axis_ids[node.get("val")])
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except Exception:
                        pass
                info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                destination.writestr(info, data)
        shutil.copy2(target, path)


def recalculate_workbook(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="offboarding-control-recalc-") as temporary:
        output_dir = Path(temporary) / "output"
        profile_dir = Path(temporary) / "profile"
        output_dir.mkdir()
        process = subprocess.run(
            ["libreoffice", "--headless", f"-env:UserInstallation=file://{profile_dir}",
             "--convert-to", "xlsx", "--outdir", str(output_dir), str(path)],
            stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300,
        )
        converted = output_dir / path.name
        if process.returncode != 0 or not converted.is_file():
            raise RuntimeError(f"LibreOffice 公式重算失败：{process.stdout[-1000:]}")
        shutil.copy2(converted, path)


def restore_freeze_panes(path: Path, sheet_count: int) -> None:
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with tempfile.TemporaryDirectory(prefix="offboarding-control-freeze-") as temporary:
        target = Path(temporary) / path.name
        sheet_names = {f"xl/worksheets/sheet{index}.xml" for index in range(1, sheet_count + 1)}
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w") as destination:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename in sheet_names:
                    root = etree.fromstring(data)
                    view = root.find(f"{{{namespace}}}sheetViews/{{{namespace}}}sheetView")
                    if view is not None:
                        for existing in view.findall(f"{{{namespace}}}pane"):
                            view.remove(existing)
                        pane = etree.Element(f"{{{namespace}}}pane")
                        pane.set("ySplit", "1")
                        pane.set("topLeftCell", "A2")
                        pane.set("activePane", "bottomLeft")
                        pane.set("state", "frozen")
                        view.insert(0, pane)
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                destination.writestr(info, data)
        shutil.copy2(target, path)


def build_workbook(departures: list[Departure], policies: list[tuple[int, str, str]], target: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.creator = "EvoCowork task-specific reference solution"
    workbook.title = "Offboarding Decision Register"
    workbook.subject = "AutomationBench 5004 policy and temporal control"
    workbook.description = "Formula-driven package built only from three frozen inputs."

    source = workbook.create_sheet("Source_Departures")
    source.append([
        "Source Row", "Employee", "Employee Email", "Department", "Manager", "Manager Email",
        "Last Day", "Type", "Existing Status", "Notes",
    ])
    for item in departures:
        source.append([
            item.source_row, item.employee, item.employee_email, item.department, item.manager,
            item.manager_email, item.last_day, item.separation_type, item.existing_status, item.notes,
        ])
        source.cell(source.max_row, 7).number_format = "yyyy-mm-dd"
    add_table(source, "SourceDepartures", "A1:J5")
    style_sheet(source, {
        "A": 11, "B": 20, "C": 31, "D": 16, "E": 20,
        "F": 31, "G": 14, "H": 29, "I": 16, "J": 39,
    }, a3=True)
    source.print_area = "A1:J5"

    policy = workbook.create_sheet("Policy_Control")
    policy.append(["Source Row", "Policy", "Steps", "Control Interpretation"])
    interpretations = {
        "Voluntary Resignation": "Public farewell allowed; IT revocation effective on last day; set Processed",
        "Involuntary Termination": "No public Slack; immediate IT revocation; confidential HR Director notice; set Processed",
        "Future-Dated Departures": "More than 14 days: Scheduled only and no notification",
        "Severance Payments": "Payroll-only after legal review; HR routes scope but cannot calculate, approve or pay",
    }
    for row_id, name, steps in policies:
        policy.append([row_id, name, steps, interpretations[name]])
    add_table(policy, "PolicyControls", "A1:D5")
    style_sheet(policy, {"A": 11, "B": 28, "C": 82, "D": 55}, a3=True)
    for row_number in range(2, 6):
        policy.row_dimensions[row_number].height = 46
    policy.print_area = "A1:D5"

    action = workbook.create_sheet("Action_Register")
    action.append([
        "Source Row", "Employee", "Employee Email", "Last Day", "Existing Status", "Type", "Notes",
        "Days From As-Of", "Target Status", "Slack Action", "IT Action", "HR Action", "Payroll Action",
        "Decision Rationale", "Source Trace",
    ])
    for row_number, item in enumerate(departures, start=2):
        action.append([
            f"=Source_Departures!A{row_number}", f"=Source_Departures!B{row_number}",
            f"=Source_Departures!C{row_number}", f"=Source_Departures!G{row_number}",
            f"=Source_Departures!I{row_number}", f"=Source_Departures!H{row_number}",
            f"=Source_Departures!J{row_number}", f"=D{row_number}-DATE(2026,3,10)",
            f'=IF(E{row_number}="Processed","NO_ACTION_ALREADY_PROCESSED",IF(H{row_number}>14,"SCHEDULED","PROCESSED"))',
            f'=IF(AND(I{row_number}="PROCESSED",LEFT(F{row_number},9)="Voluntary"),"DRAFT_FAREWELL","NO_PUBLIC_POST")',
            f'=IF(I{row_number}<>"PROCESSED","NO_ACTION",IF(ISNUMBER(SEARCH("Involuntary",F{row_number})),"DRAFT_REVOKE_IMMEDIATE","DRAFT_REVOKE_LAST_DAY"))',
            f'=IF(AND(I{row_number}="PROCESSED",ISNUMBER(SEARCH("Involuntary",F{row_number}))),"DRAFT_HR_NOTICE","NO_ACTION")',
            f'=IF(I{row_number}="PROCESSED","ROUTE_SCOPE_REVIEW","NO_ACTION")',
            item.rationale,
            f"Employee_Departures.xlsx!Departures row {item.source_row}",
        ])
        action.cell(row_number, 4).number_format = "yyyy-mm-dd"
    add_table(action, "OffboardingActions", "A1:O5")
    style_sheet(action, {
        "A": 7, "B": 14, "C": 20, "D": 11, "E": 12, "F": 18, "G": 22, "H": 10,
        "I": 18, "J": 15, "K": 18, "L": 15, "M": 17, "N": 30, "O": 24,
    }, a3=True)
    for cell in action[1]:
        cell.font = Font(name="Aptos", bold=True, color=WHITE, size=8)
    for row_number in range(2, 6):
        action.row_dimensions[row_number].height = 50
    action.conditional_formatting.add("A2:O5", FormulaRule(formula=['$I2="PROCESSED"'], fill=PatternFill("solid", fgColor=GREEN)))
    action.conditional_formatting.add("A2:O5", FormulaRule(formula=['$I2="SCHEDULED"'], fill=PatternFill("solid", fgColor=AMBER)))
    action.conditional_formatting.add("A2:O5", FormulaRule(formula=['$I2="NO_ACTION_ALREADY_PROCESSED"'], fill=PatternFill("solid", fgColor=RED)))
    action.print_area = "A1:O5"

    ledger = workbook.create_sheet("Notification_Ledger")
    ledger.append([
        "Channel", "Action Type", "Employee Scope", "Recipient / Channel", "Subject / Title",
        "Message ID", "Delivery Status", "Effective Date", "Control Note", "Source Rows",
    ])
    rows = [
        ("EMAIL", "IT_REVOKE_LAST_DAY", "Greg Foster", "it-ops@company.example.com", "Access revocation request — Greg Foster", "<offboarding-it-row2@evocowork.invalid>", "DRAFT_NOT_SENT", "2026-03-07", "Revoke for last day; not executed", "2"),
        ("EMAIL", "IT_REVOKE_IMMEDIATE", "Diana Reese", "it-ops@company.example.com", "Confidential immediate access revocation — Diana Reese", "<offboarding-it-row3@evocowork.invalid>", "DRAFT_NOT_SENT", "2026-03-10", "Immediate request; not executed", "3"),
        ("EMAIL", "HR_DIRECTOR_NOTICE", "Diana Reese", "hr-director@company.example.com", "Confidential separation notice — Diana Reese", "<offboarding-hr-row3@evocowork.invalid>", "DRAFT_NOT_SENT", "2026-03-10", "Legal review note preserved; no public post", "3"),
        ("EMAIL", "PAYROLL_SCOPE_REVIEW", "Greg Foster; Diana Reese", "payroll@company.example.com", "Payroll-owned severance scope review — current offboarding cases", "<offboarding-payroll-batch@evocowork.invalid>", "DRAFT_NOT_SENT", "2026-03-10", "Review only; no amount, approval or payment", "2;3"),
        ("SLACK", "PUBLIC_FAREWELL", "Greg Foster", "C_DEPART / departures", "Farewell draft — Greg Foster", "offboarding-slack-row2", "DRAFT_NOT_POSTED", "2026-03-10", "Only voluntary current case; not posted", "2"),
    ]
    for row in rows:
        ledger.append(row)
    add_table(ledger, "NotificationLedger", "A1:J6")
    style_sheet(ledger, {"A": 8, "B": 17, "C": 18, "D": 24, "E": 31, "F": 27, "G": 15, "H": 12, "I": 30, "J": 8}, a3=True)
    for cell in ledger[1]:
        cell.font = Font(name="Aptos", bold=True, color=WHITE, size=8)
    for row_number in range(2, 7):
        ledger.row_dimensions[row_number].height = 42
    ledger.print_area = "A1:J6"

    summary = workbook.create_sheet("Control_Summary")
    summary.append(["Metric", "Formula Result", "Expected", "Control Status", "Evidence"])
    metrics = [
        ("Source departures", "=COUNTA(Source_Departures!B2:B5)", 4, "Four frozen source rows"),
        ("Process now", '=COUNTIF(Action_Register!I2:I5,"PROCESSED")', 2, "Greg Foster + Diana Reese"),
        ("Scheduled only", '=COUNTIF(Action_Register!I2:I5,"SCHEDULED")', 1, "Raj Patel at +36 days"),
        ("Already processed / no action", '=COUNTIF(Action_Register!I2:I5,"NO_ACTION_ALREADY_PROCESSED")', 1, "Greg Forster source status"),
        ("Public Slack drafts", '=COUNTIF(Action_Register!J2:J5,"DRAFT_FAREWELL")', 1, "Greg Foster only"),
        ("IT drafts", '=COUNTIF(Action_Register!K2:K5,"DRAFT_REVOKE*")', 2, "Rows 2 and 3"),
        ("HR Director drafts", '=COUNTIF(Action_Register!L2:L5,"DRAFT_HR_NOTICE")', 1, "Diana Reese only"),
        ("Payroll scope reviews", '=COUNTIF(Action_Register!M2:M5,"ROUTE_SCOPE_REVIEW")/2', 1, "One combined draft for two actionable cases"),
        ("Email drafts", '=COUNTIF(Notification_Ledger!A2:A6,"EMAIL")', 4, "Four MBOX messages"),
        ("Slack drafts", '=COUNTIF(Notification_Ledger!A2:A6,"SLACK")', 1, "One channel draft"),
    ]
    for row_number, (metric, formula, expected, evidence) in enumerate(metrics, start=2):
        summary.append([metric, formula, expected, f'=IF(B{row_number}=C{row_number},"PASS","FAIL")', evidence])
    add_table(summary, "OffboardingMetrics", "A1:E11")
    style_sheet(summary, {"A": 34, "B": 22, "C": 18, "D": 20, "E": 48})
    summary.conditional_formatting.add("A2:E11", FormulaRule(formula=['$D2="PASS"'], fill=PatternFill("solid", fgColor=GREEN)))
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Offboarding Decision Outcomes"
    chart.y_axis.title = "Records"
    chart.x_axis.title = "Outcome"
    chart.height = 7.2
    chart.width = 13.0
    chart.legend = None
    chart.add_data(Reference(summary, min_col=2, min_row=3, max_row=5), titles_from_data=False)
    chart.set_categories(Reference(summary, min_col=1, min_row=3, max_row=5))
    summary.add_chart(chart, "G2")
    summary.print_area = "A1:N18"

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(target)
    recalculate_workbook(target)
    restore_freeze_panes(target, 5)
    normalize_ooxml(target)


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    properties.append(shade)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = "0F172A", size: int = 7) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph()
    paragraph.style = document.styles[f"Heading {level}"]
    run = paragraph.add_run(text)
    run.font.name = "Aptos Display"
    run.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else BLUE)


def build_docx(departures: list[Departure], target: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.25)
    section.bottom_margin = Cm(1.25)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9)
    for name, size in (("Heading 1", 18), ("Heading 2", 12)):
        styles[name].font.name = "Aptos Display"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True

    title = document.add_paragraph()
    run = title.add_run("Offboarding Control Brief")
    run.font.name = "Aptos Display"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = document.add_paragraph("AutomationBench 5004 | Offline control package | As of 2026-03-10")
    subtitle.runs[0].font.color.rgb = RGBColor.from_string("475569")

    callout = document.add_table(rows=1, cols=4)
    callout.style = "Table Grid"
    callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    cards = [("2", "process now"), ("1", "scheduled"), ("1", "already processed"), ("5", "offline drafts")]
    for cell, (value, label) in zip(callout.rows[0].cells, cards):
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        value_run = paragraph.add_run(value + "\n")
        value_run.bold = True
        value_run.font.name = "Aptos Display"
        value_run.font.size = Pt(16)
        value_run.font.color.rgb = RGBColor.from_string(NAVY)
        label_run = paragraph.add_run(label)
        label_run.font.name = "Aptos"
        label_run.font.size = Pt(8)
        label_run.font.color.rgb = RGBColor.from_string("475569")
        shade_cell(cell, LIGHT_BLUE)

    add_heading(document, "Decision register", 1)
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Employee", "Last day", "Δ days", "Source", "Outcome", "Action plan"]
    widths = [3.0, 2.1, 1.3, 2.0, 2.8, 6.2]
    for cell, header, width in zip(table.rows[0].cells, headers, widths):
        cell.width = Cm(width)
        set_cell_text(cell, header, bold=True, color=WHITE, size=7)
        shade_cell(cell, NAVY)
    for item in departures:
        row = table.add_row().cells
        if item.employee == "Greg Foster":
            action_plan = "Slack farewell; IT last-day; Payroll review"
        elif item.employee == "Diana Reese":
            action_plan = "No public Slack; IT immediate; HR notice; Payroll review"
        elif item.employee == "Raj Patel":
            action_plan = "Schedule only; no notification"
        else:
            action_plan = "Already processed; no repeat action"
        outcome = (
            "PROCESSED" if item.target_status == "PROCESSED"
            else "SCHEDULED" if item.target_status == "SCHEDULED"
            else "NO ACTION"
        )
        values = [
            f"{item.employee}\nrow {item.source_row}", item.last_day.isoformat(), str(item.days_from_as_of),
            item.existing_status, outcome, action_plan,
        ]
        for cell, value, width in zip(row, values, widths):
            cell.width = Cm(width)
            set_cell_text(cell, value, size=6)
        shade_cell(row[4], GREEN if item.target_status == "PROCESSED" else AMBER if item.target_status == "SCHEDULED" else RED)

    add_heading(document, "Prepared drafts", 1)
    draft_lines = [
        "DRAFT_REVOKE_LAST_DAY — IT / Greg Foster — 2026-03-07 — <offboarding-it-row2@evocowork.invalid>",
        "DRAFT_REVOKE_IMMEDIATE — IT / Diana Reese — <offboarding-it-row3@evocowork.invalid>",
        "DRAFT_HR_NOTICE — HR Director / Diana Reese — <offboarding-hr-row3@evocowork.invalid>",
        "ROUTE_SCOPE_REVIEW — Payroll / Greg Foster + Diana Reese — <offboarding-payroll-batch@evocowork.invalid>",
        "DRAFT_FAREWELL — #departures / Greg Foster — offboarding-slack-row2",
    ]
    for line in draft_lines:
        document.add_paragraph(line, style="List Bullet")
    boundary = document.add_paragraph()
    run = boundary.add_run("CONTROL BOUNDARY  ")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("9C5700")
    boundary.add_run("All email is DRAFT_NOT_SENT and Slack is DRAFT_NOT_POSTED. No revocation or payment has been executed.")
    document.add_paragraph("Diana Reese is NO_PUBLIC_POST under the involuntary-separation confidentiality rule.")

    document.add_page_break()
    add_heading(document, "Policy precedence and exception audit", 1)
    document.add_paragraph(
        "The source policy controls the user request. Existing Processed rows are never replayed. A pending last day more than 14 days away is Scheduled with no notification. Involuntary departures stay out of public Slack. Severance remains Payroll-only after legal review; HR Ops can route scope for review but cannot calculate, approve, authorize or pay it."
    )
    exception_table = document.add_table(rows=1, cols=4)
    exception_table.style = "Table Grid"
    exception_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, header in zip(exception_table.rows[0].cells, ["Control", "Entity", "Evidence", "Outcome"]):
        set_cell_text(cell, header, bold=True, color=WHITE, size=8)
        shade_cell(cell, NAVY)
    exceptions = [
        ("Confidentiality", "Diana Reese", "Involuntary - Position Eliminated", "No public Slack; IT + HR Director drafts"),
        ("Legal review", "Diana Reese", "HR and legal reviewing this case", "Payroll review only; no payment authorization"),
        ("14-day boundary", "Raj Patel", "2026-04-15 = +36 days", "Scheduled; no notification"),
        ("Near-name decoy", "Greg Forster", "Source Status = Processed", "NO_ACTION_ALREADY_PROCESSED; distinct from Greg Foster"),
        ("Scope conflict", "Severance request", "Payroll-only policy", "HR does not calculate, approve or pay"),
    ]
    for values in exceptions:
        row = exception_table.add_row().cells
        for cell, value in zip(row, values):
            set_cell_text(cell, value, size=7)

    add_heading(document, "Traceability map", 1)
    trace = document.add_table(rows=1, cols=3)
    trace.style = "Table Grid"
    trace.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, header in zip(trace.rows[0].cells, ["Input", "Decision role", "Output trace"]):
        set_cell_text(cell, header, bold=True, color=WHITE, size=8)
        shade_cell(cell, NAVY)
    trace_rows = [
        ("Employee_Departures.xlsx", "Four source rows + four controlling policies", "Source_Departures, Policy_Control, Action_Register"),
        ("Workflow_Context.json", "As-of time 2026-03-10T09:00:00Z", "Days From As-Of and target status"),
        ("Departures_Channel.json", "C_DEPART / departures", "Departures_Channel_Draft.json"),
    ]
    for values in trace_rows:
        row = trace.add_row().cells
        for cell, value in zip(row, values):
            set_cell_text(cell, value, size=7)

    add_heading(document, "Final control statement", 1)
    paragraph = document.add_paragraph()
    paragraph.add_run("Prepared, not executed: ").bold = True
    paragraph.add_run("four email drafts and one public farewell draft. Greg Foster and Diana Reese enter the current Processed workflow; Raj Patel is Scheduled; Greg Forster remains already processed with no action.")
    paragraph = document.add_paragraph()
    paragraph.add_run("Severance boundary: ").bold = True
    paragraph.add_run("No severance amount was calculated or approved, no payment was authorized or made, and Diana Reese remains pending legal review.")

    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_run = footer.add_run("EvoCowork | AutomationBench 5004 | Offline draft controls")
        footer_run.font.name = "Aptos"
        footer_run.font.size = Pt(7)
        footer_run.font.color.rgb = RGBColor.from_string("64748B")
    document.save(target)
    normalize_ooxml(target)


def email_bytes(*, to: str, subject: str, message_id: str, source_rows: str, action: str, body: str) -> bytes:
    message = EmailMessage(policy=SMTP)
    message["From"] = "hr-ops@company.example.com"
    message["To"] = to
    message["Subject"] = subject
    message["Date"] = "Tue, 10 Mar 2026 09:00:00 +0000"
    message["Message-ID"] = message_id
    message["X-EvoCowork-Delivery-Status"] = "DRAFT_NOT_SENT"
    message["X-EvoCowork-Source-Rows"] = source_rows
    message["X-EvoCowork-Action-Type"] = action
    message.set_content("Draft status: DRAFT_NOT_SENT. This message has not been sent.\n\n" + body)
    buffer = BytesIO()
    BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
    return buffer.getvalue()


def build_mbox(target: Path) -> None:
    messages = [
        email_bytes(
            to="it-ops@company.example.com", subject="Access revocation request — Greg Foster",
            message_id="<offboarding-it-row2@evocowork.invalid>", source_rows="2", action="IT_REVOKE_LAST_DAY",
            body=("Please prepare access revocation for Greg Foster (greg.foster@company.example.com), Sales, manager Lisa Wang. "
                  "Last day: 2026-03-07. Separation type: Voluntary - Resignation. Execute only after authorized review; this draft does not prove revocation."),
        ),
        email_bytes(
            to="it-ops@company.example.com", subject="Confidential immediate access revocation — Diana Reese",
            message_id="<offboarding-it-row3@evocowork.invalid>", source_rows="3", action="IT_REVOKE_IMMEDIATE",
            body=("CONFIDENTIAL. Please prepare immediate access revocation for Diana Reese (diana.reese@company.example.com), Engineering, manager David Okonkwo. "
                  "Last day: 2026-03-14. Separation type: Involuntary - Position Eliminated. This request remains an offline draft and has not been executed."),
        ),
        email_bytes(
            to="hr-director@company.example.com", subject="Confidential separation notice — Diana Reese",
            message_id="<offboarding-hr-row3@evocowork.invalid>", source_rows="3", action="HR_DIRECTOR_NOTICE",
            body=("CONFIDENTIAL. Diana Reese, Engineering, manager David Okonkwo, has a 2026-03-14 last day under Involuntary - Position Eliminated. "
                  "Source note preserved verbatim: HR and legal reviewing this case. No public Slack announcement has been prepared. This notice has not been sent."),
        ),
        email_bytes(
            to="payroll@company.example.com", subject="Payroll-owned severance scope review — current offboarding cases",
            message_id="<offboarding-payroll-batch@evocowork.invalid>", source_rows="2;3", action="PAYROLL_SCOPE_REVIEW",
            body=("ROUTE_SCOPE_REVIEW_ONLY. Please review severance scope under the Payroll-only policy for Greg Foster (last day 2026-03-07) and Diana Reese (last day 2026-03-14). "
                  "Diana Reese source note: HR and legal reviewing this case. No severance amount was calculated or approved. No payment has been authorized or made. HR Ops is not processing severance."),
        ),
    ]
    chunks: list[bytes] = []
    for payload in messages:
        chunks.append(b"From hr-ops@company.example.com Tue Mar 10 09:00:00 2026\n")
        chunks.append(payload.replace(b"\r\n", b"\n"))
        chunks.append(b"\n")
    target.write_bytes(b"".join(chunks))


def build_slack(channel: dict, target: Path) -> None:
    channel_record = channel["channels"][0]
    payload = {
        "task_id": "cowork-evo/automationbench-5004-offboarding-control-package",
        "channel": {"id": channel_record["id"], "name": channel_record["name"]},
        "message_id": "offboarding-slack-row2",
        "status": "DRAFT_NOT_POSTED",
        "source_row": 2,
        "employee": "Greg Foster",
        "linked_message_ids": ["<offboarding-it-row2@evocowork.invalid>"],
        "source_files": ["Employee_Departures.xlsx", "Workflow_Context.json", "Departures_Channel.json"],
        "text": (
            "DRAFT_NOT_POSTED — Farewell draft for Greg Foster, Sales. Greg's last day was 2026-03-07. "
            "Thank you for your contributions, Greg, and best wishes for what comes next. This message has not been posted."
        ),
    }
    target.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    departures, policies, channel = read_inputs(args.input_dir)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for existing in args.output_dir.iterdir():
        if existing.is_file() or existing.is_symlink():
            existing.unlink()
        elif existing.is_dir():
            shutil.rmtree(existing)
    build_workbook(departures, policies, args.output_dir / XLSX_NAME)
    build_docx(departures, args.output_dir / DOCX_NAME)
    build_mbox(args.output_dir / MBOX_NAME)
    build_slack(channel, args.output_dir / SLACK_NAME)


if __name__ == "__main__":
    main()
