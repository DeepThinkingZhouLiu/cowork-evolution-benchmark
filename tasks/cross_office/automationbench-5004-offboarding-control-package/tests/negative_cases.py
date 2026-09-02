#!/usr/bin/env python3
"""AutomationBench 5004 的逐项定向负例回放。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


XLSX = "Offboarding_Decision_Register.xlsx"
DOCX = "Offboarding_Control_Brief.docx"
MBOX = "Offboarding_Communications.mbox"
SLACK = "Departures_Channel_Draft.json"


def save_workbook(path: Path, mutate) -> None:
    workbook = load_workbook(path)
    mutate(workbook)
    workbook.save(path)


def replace_docx(path: Path, old: str, new: str) -> None:
    document = Document(path)
    for paragraph in document.paragraphs:
        if old in paragraph.text:
            for run in paragraph.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
            if old in paragraph.text:
                paragraph.text = paragraph.text.replace(old, new)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if old in cell.text:
                    cell.text = cell.text.replace(old, new)
    document.save(path)


def write_messages(path: Path, messages: list) -> None:
    chunks: list[bytes] = []
    for message in messages:
        chunks.append(b"From hr-ops@company.example.com Tue Mar 10 09:00:00 2026\n")
        buffer = BytesIO()
        BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
        chunks.append(buffer.getvalue().replace(b"\r\n", b"\n"))
        chunks.append(b"\n")
    path.write_bytes(b"".join(chunks))


def mutate_message_header(path: Path, message_id: str, header: str, value: str) -> None:
    messages = list(mailbox.mbox(path, create=False))
    for message in messages:
        if str(message.get("Message-ID", "")) == message_id:
            if message.get(header) is None:
                message[header] = value
            else:
                message.replace_header(header, value)
    write_messages(path, messages)


def raw_replace(path: Path, old: bytes, new: bytes) -> None:
    data = path.read_bytes()
    if old not in data:
        raise ValueError(f"负例目标字节不存在：{old!r}")
    path.write_bytes(data.replace(old, new, 1))


def replace_message_body(path: Path, message_id: str, old: str, new: str) -> None:
    messages = list(mailbox.mbox(path, create=False))
    rewritten = []
    replaced = False
    for message in messages:
        if str(message.get("Message-ID", "")) != message_id:
            rewritten.append(message)
            continue
        payload = message.get_payload(decode=True)
        body = payload.decode(message.get_content_charset() or "utf-8", errors="replace") if payload else str(message.get_payload())
        if old not in body:
            raise ValueError(f"负例目标正文不存在：{old!r}")
        replacement = EmailMessage(policy=SMTP)
        for key, value in message.items():
            if key.lower() not in {"content-type", "content-transfer-encoding", "mime-version"}:
                replacement[key] = value
        replacement.set_content(body.replace(old, new, 1), charset="utf-8")
        rewritten.append(replacement)
        replaced = True
    if not replaced:
        raise ValueError(f"负例目标 Message-ID 不存在：{message_id}")
    write_messages(path, rewritten)


def cases() -> dict[str, callable]:
    def missing_docx(root: Path) -> None:
        (root / DOCX).unlink()

    def extra_debug(root: Path) -> None:
        (root / "debug.txt").write_text("unexpected", encoding="utf-8")

    def source_name_changed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Source_Departures"]["B5"], "value", "Greg Foster"))

    def source_date_changed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Source_Departures"]["G4"], "value", "2026-03-20"))

    def policy_severance_altered(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Policy_Control"]["C5"], "value", "HR may process severance directly."))

    def greg_foster_no_slack(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Action_Register"]["J2"], "value", "NO_PUBLIC_POST"))

    def diana_public_slack(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Action_Register"]["J3"], "value", "DRAFT_FAREWELL"))

    def diana_it_last_day(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Action_Register"]["K3"], "value", "DRAFT_REVOKE_LAST_DAY"))

    def diana_note_removed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Action_Register"]["G3"], "value", ""))

    def raj_processed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Action_Register"]["I4"], "value", "PROCESSED"))

    def forster_reprocessed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Action_Register"]["I5"], "value", "PROCESSED"))

    def name_confusion(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Action_Register"]["B5"], "value", "Greg Foster"))

    def days_staticized(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Action_Register"]["H4"], "value", 36))

    def status_staticized(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Action_Register"]["I2"], "value", "PROCESSED"))

    def summary_wrong(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Control_Summary"]["C3"], "value", 3))

    def native_chart_removed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: wb["Control_Summary"]._charts.clear())

    def native_tables_removed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: wb["Action_Register"].tables.clear())

    def corrupt_xlsx(root: Path) -> None:
        (root / XLSX).write_bytes(b"not an xlsx")

    def mbox_drop_message(root: Path) -> None:
        messages = list(mailbox.mbox(root / MBOX, create=False))[1:]
        write_messages(root / MBOX, messages)

    def mbox_employee_recipient(root: Path) -> None:
        mutate_message_header(root / MBOX, "<offboarding-it-row2@evocowork.invalid>", "To", "greg.foster@company.example.com")

    def mbox_raj_added(root: Path) -> None:
        raw_replace(root / MBOX, b"Greg Foster (last day 2026-03-07) and Diana Reese", b"Greg Foster (last day 2026-03-07), Raj Patel, and Diana Reese")

    def mbox_payroll_amount(root: Path) -> None:
        replace_message_body(root / MBOX, "<offboarding-payroll-batch@evocowork.invalid>", "No severance amount was calculated or approved.", "A $25,000 severance amount was calculated and approved.")

    def mbox_false_paid(root: Path) -> None:
        replace_message_body(root / MBOX, "<offboarding-payroll-batch@evocowork.invalid>", "No payment has been authorized or made.", "Severance has been paid.")

    def mbox_false_sent(root: Path) -> None:
        replace_message_body(root / MBOX, "<offboarding-it-row2@evocowork.invalid>", "This message has not been sent.", "This message was sent successfully.")

    def mbox_message_id_conflict(root: Path) -> None:
        raw_replace(root / MBOX, b"<offboarding-it-row3@evocowork.invalid>", b"<offboarding-it-row2@evocowork.invalid>")

    def slack_add_diana(root: Path) -> None:
        payload = json.loads((root / SLACK).read_text(encoding="utf-8"))
        payload["text"] += " Farewell Diana Reese."
        (root / SLACK).write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    def slack_wrong_person(root: Path) -> None:
        payload = json.loads((root / SLACK).read_text(encoding="utf-8"))
        payload["employee"] = "Greg Forster"
        payload["source_row"] = 5
        (root / SLACK).write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    def slack_false_posted(root: Path) -> None:
        payload = json.loads((root / SLACK).read_text(encoding="utf-8"))
        payload["status"] = "POSTED"
        payload["text"] = payload["text"].replace("DRAFT_NOT_POSTED", "POSTED").replace("has not been posted", "was posted to #departures")
        (root / SLACK).write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    def slack_message_id_conflict(root: Path) -> None:
        payload = json.loads((root / SLACK).read_text(encoding="utf-8"))
        payload["message_id"] = "offboarding-slack-row3"
        (root / SLACK).write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    def word_wrong_days(root: Path) -> None:
        replace_docx(root / DOCX, "2026-04-15 = +36 days", "2026-04-15 = +14 days")

    def word_remove_policy_boundary(root: Path) -> None:
        replace_docx(root / DOCX, "Payroll-only", "HR-owned")

    def word_false_executed(root: Path) -> None:
        replace_docx(root / DOCX, "No revocation or payment has been executed.", "Access was revoked and payment was authorized.")

    return {
        "missing-docx": missing_docx,
        "extra-debug-file": extra_debug,
        "source-name-changed": source_name_changed,
        "source-date-changed": source_date_changed,
        "policy-severance-altered": policy_severance_altered,
        "greg-foster-no-slack": greg_foster_no_slack,
        "diana-public-slack-substring-regression": diana_public_slack,
        "diana-it-last-day": diana_it_last_day,
        "diana-legal-note-removed": diana_note_removed,
        "raj-marked-processed": raj_processed,
        "forster-reprocessed": forster_reprocessed,
        "near-name-confusion": name_confusion,
        "days-staticized": days_staticized,
        "status-staticized": status_staticized,
        "summary-wrong": summary_wrong,
        "native-chart-removed": native_chart_removed,
        "native-tables-removed": native_tables_removed,
        "corrupt-xlsx": corrupt_xlsx,
        "mbox-drop-message": mbox_drop_message,
        "mbox-employee-recipient": mbox_employee_recipient,
        "mbox-add-scheduled-raj": mbox_raj_added,
        "mbox-payroll-add-amount": mbox_payroll_amount,
        "mbox-false-paid": mbox_false_paid,
        "mbox-false-sent": mbox_false_sent,
        "mbox-message-id-conflict": mbox_message_id_conflict,
        "slack-add-confidential-diana": slack_add_diana,
        "slack-wrong-greg": slack_wrong_person,
        "slack-false-posted": slack_false_posted,
        "slack-message-id-conflict": slack_message_id_conflict,
        "word-wrong-days": word_wrong_days,
        "word-remove-payroll-boundary": word_remove_policy_boundary,
        "word-false-executed": word_false_executed,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--golden", type=Path, default=Path(__file__).resolve().parents[1] / "solution" / "Golden")
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()

    judge_path = Path(__file__).with_name("judge.py")
    spec = importlib.util.spec_from_file_location("offboarding_task_judge", judge_path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入本题 Judge")
    judge = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = judge
    spec.loader.exec_module(judge)

    results: list[dict] = []
    with tempfile.TemporaryDirectory(prefix="offboarding-negative-cases-") as temporary:
        root = Path(temporary)
        for name, mutate in cases().items():
            case_dir = root / name
            shutil.copytree(args.golden, case_dir)
            mutate(case_dir)
            result = judge.evaluate(case_dir)
            results.append({
                "case": name,
                "reward": result["reward"],
                "passed": result["passed"],
                "hurdle_failures": result.get("hurdle_failures", []),
                "penalties": [item["id"] for item in result.get("penalties", [])],
            })

    failed = [item for item in results if item["passed"] or float(item["reward"]) >= 0.8]
    report = {
        "total": len(results),
        "rejected": len(results) - len(failed),
        "maximum_reward": max(float(item["reward"]) for item in results),
        "cases": results,
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if failed:
        raise SystemExit(f"存在误接受负例：{[item['case'] for item in failed]}")


if __name__ == "__main__":
    main()
