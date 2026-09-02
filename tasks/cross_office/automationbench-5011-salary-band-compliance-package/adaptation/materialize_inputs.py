#!/usr/bin/env python3
"""只静态化 AutomationBench 5011；不是通用造题器。"""
import hashlib,importlib.util,json,shutil,sys,types
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.worksheet.table import Table,TableStyleInfo
ROOT=Path(__file__).resolve().parents[1];UP=Path(__file__).with_name("upstream-hr-tasks.py");SNAP=Path(__file__).with_name("upstream-task-5011.json");DATA=ROOT/"data/input_files";ENV=ROOT/"environment/input_files";UPSTREAM_SHA="c173a36a03f43096bd122577584aa32f9df379d8760aa46c43be04c85d101779"
def load_task():
 if hashlib.sha256(UP.read_bytes()).hexdigest()!=UPSTREAM_SHA:raise ValueError("上游哈希错误")
 d=types.ModuleType("datasets");d.Dataset=object;sys.modules["datasets"]=d
 for name in ("automationbench","automationbench.domains","automationbench.domains.hr"):sys.modules[name]=types.ModuleType(name)
 noise=types.ModuleType("automationbench.domains.hr._noise");noise.apply_noise=lambda x:x;sys.modules[noise.__name__]=noise
 spec=importlib.util.spec_from_file_location("automationbench_5011",UP);m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m);task=m.get_hr_salary_band_audit_task();assert task["example_id"]==5011 and task["task"]=="hr.salary_band_audit";return task
def audit(task):
 info=task["info"];state=info["initial_state"];assert state["meta"]["current_time"]=="2026-03-15T09:00:00Z" and len(info["assertions"])==8
 sheets=state["google_sheets"]["spreadsheets"][0]["worksheets"];assert [s["title"] for s in sheets]==["Salary Bands","Employee Compensation","Audit Flag Convention"] and [len(s["rows"]) for s in sheets]==[5,6,2]
 assert [r["cells"]["Level"] for r in sheets[0]["rows"]]==["IC1","IC2","IC3","IC4","M1"] and [r["cells"]["Employee"] for r in sheets[1]["rows"]]==["Alice Park","Bob Chen","Carol Diaz","Dave Kim","Eve Liu","Lisa Wang"]
 assert [r["cells"]["Salary"] for r in sheets[1]["rows"]]==["155000","125000","108000","72000","215000","148000"]
def style(ws,widths,name):
 for i,w in enumerate(widths,1):ws.column_dimensions[ws.cell(1,i).column_letter].width=w
 ws.freeze_panes="A2";ws.sheet_view.showGridLines=False
 for c in ws[1]:c.fill=PatternFill("solid",fgColor="17365D");c.font=Font(name="Carlito",bold=True,color="FFFFFF");c.alignment=Alignment(horizontal="center",wrap_text=True)
 for row in ws.iter_rows(min_row=2):
  for c in row:c.font=Font(name="Carlito",size=9);c.alignment=Alignment(vertical="top",wrap_text=True)
 t=Table(displayName=name,ref=ws.dimensions);t.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True);ws.add_table(t);ws.page_setup.orientation="landscape";ws.page_setup.fitToWidth=1;ws.sheet_properties.pageSetUpPr.fitToPage=True
def write_xlsx(sheets,path):
 wb=Workbook();wb.remove(wb.active)
 configs=[("Salary Bands",["Source Row","Level","Title Pattern","Min","Mid","Max"],[10,16,32,18,18,18],"SalaryBandsSource"),("Employee Compensation",["Source Row","Employee","Title","Level","Salary","Flag"],[10,24,34,16,18,16],"EmployeeCompSource"),("Audit Flag Convention",["Source Row","Condition","Flag"],[10,52,18],"FlagConventionSource")]
 for source,(title,headers,widths,name) in zip(sheets,configs):
  ws=wb.create_sheet(title);ws.append(headers)
  for r in source["rows"]:
   if title=="Salary Bands":vals=[r["row_id"],r["cells"]["Level"],r["cells"]["Title Pattern"],int(r["cells"]["Min"]),int(r["cells"]["Mid"]),int(r["cells"]["Max"])]
   elif title=="Employee Compensation":vals=[r["row_id"],r["cells"]["Employee"],r["cells"]["Title"],r["cells"]["Level"],int(r["cells"]["Salary"]),r["cells"]["Flag"]]
   else:vals=[r["row_id"],r["cells"]["Condition"],r["cells"]["Flag"]]
   ws.append(vals)
  style(ws,widths,name)
  if title!="Audit Flag Convention":
   for row in ws.iter_rows(min_row=2):
    for c in row:
     if c.column in ({4,5,6} if title=="Salary Bands" else {5}):c.number_format='0'
 wb.save(path)
def main():
 task=load_task();audit(task);state=task["info"]["initial_state"];DATA.mkdir(parents=True,exist_ok=True);ENV.mkdir(parents=True,exist_ok=True);snapshot=json.dumps(task,ensure_ascii=False,indent=2)+"\n"
 if SNAP.exists() and SNAP.read_text()!=snapshot:raise ValueError("单题快照漂移")
 SNAP.write_text(snapshot);write_xlsx(state["google_sheets"]["spreadsheets"][0]["worksheets"],DATA/"Salary_Band_Audit_Source.xlsx");(DATA/"Workflow_Context.json").write_text(json.dumps({"meta":state["meta"],"report_recipient":"comp-team@company.example.com"},indent=2)+"\n")
 for item in DATA.iterdir():
  if item.is_file():shutil.copy2(item,ENV/item.name)
if __name__=="__main__":main()
