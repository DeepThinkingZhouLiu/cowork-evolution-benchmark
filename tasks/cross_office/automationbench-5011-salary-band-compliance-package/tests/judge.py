#!/usr/bin/env python3
"""AutomationBench 5011 的 v2 原子部分分 Judge；v1 保存在 judge_v1.py。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import mailbox
import re
import sys
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from office_checks import body_text, package_safe, render


TASK_ID = "cowork-evo/automationbench-5011-salary-band-compliance-package"
RUBRIC_VERSION = "automationbench-5011-v0.2"
JUDGE_VERSION = "automationbench-5011-atomic-partial-v2.0"
X = "Salary_Band_Compliance_Audit.xlsx"
D = "Salary_Band_Compliance_Brief.docx"
P = "Salary_Band_Compliance_Review.pptx"
M = "Compensation_Team_Report_Draft.mbox"
U = "Salary_Flag_Update_Drafts.json"
C = "Salary_Band_Execution_Control.json"
REQUIRED = (X, D, P, M, U, C)
WEIGHTS = {"R001":6,"R002":8,"R003":10,"R004":10,"R005":10,"R006":8,"R007":9,"R008":8,"R009":8,"R010":8,"R011":7,"R012":8,"R013":-15,"R014":-15}
EXPECTED_SHEETS = ["Source_Bands","Source_Employees","Source_Flag_Convention","Evidence_Register","Band_Check","Flagged_Employees","Within_Band","Flag_Update_Drafts","Report_Draft","Control_Summary"]
EXPECTED_BANDS = [
    (2,"IC1","Junior/Associate",55000,70000,85000),
    (3,"IC2","Mid-level (no prefix)",80000,100000,120000),
    (4,"IC3","Senior",110000,135000,160000),
    (5,"IC4","Staff/Principal",150000,180000,210000),
    (6,"M1","Manager",120000,145000,170000),
]
EXPECTED_EMPLOYEES = [
    (2,"Alice Park","Senior Software Engineer","IC3",155000,None),
    (3,"Bob Chen","Software Engineer","IC2",125000,None),
    (4,"Carol Diaz","Senior Marketing Analyst","IC3",108000,None),
    (5,"Dave Kim","Junior Developer","IC1",72000,None),
    (6,"Eve Liu","Staff Data Scientist","IC4",215000,None),
    (7,"Lisa Wang","Sales Manager","M1",148000,None),
]
EXPECTED_CONVENTIONS = [(2,"Salary above band maximum","Over"),(3,"Salary below band minimum","Under")]
EXPECTED_FLAGGED = [
    ("Bob Chen","IC2",125000,80000,120000,"Over",5000),
    ("Carol Diaz","IC3",108000,110000,160000,"Under",-2000),
    ("Eve Liu","IC4",215000,150000,210000,"Over",5000),
]
EXPECTED_WITHIN = [
    ("Alice Park","IC3",155000,110000,160000,"Within"),
    ("Dave Kim","IC1",72000,55000,85000,"Within"),
    ("Lisa Wang","M1",148000,120000,170000,"Within"),
]


def _load_v1():
    path = Path(__file__).with_name("judge_v1.py")
    spec = importlib.util.spec_from_file_location("automationbench_5011_judge_v1", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v1 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


V1 = _load_v1()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


class Units:
    def __init__(self) -> None:
        self.values: dict[str, dict[str, Any]] = {
            key: {"earned": 0.0, "total": 0.0, "failed": []}
            for key in WEIGHTS
        }

    def check(self, criterion: str, condition: bool, label: str, count: float = 1.0) -> None:
        item = self.values[criterion]
        item["total"] += count
        if condition:
            item["earned"] += count
        else:
            item["failed"].append(label)

    def violation(self, criterion: str, present: bool, label: str) -> None:
        item = self.values[criterion]
        item["total"] += 1.0
        if present:
            item["earned"] += 1.0
            item["failed"].append(label)

    def score(self, criterion: str) -> float:
        item = self.values[criterion]
        return round(item["earned"] / item["total"], 6) if item["total"] else 0.0

    def result(self, criterion: str) -> dict[str, Any]:
        item = self.values[criterion]
        failures = item["failed"]
        if WEIGHTS[criterion] < 0:
            evidence = "发现违规：" + "；".join(failures) if failures else "未发现本项违规"
        else:
            evidence = "未通过：" + "；".join(failures) if failures else f"全部 {int(item['total'])} 个任务特定检查单位通过"
        return {
            "criterion_id": criterion,
            "score": self.score(criterion),
            "evidence": evidence,
            "method": "deterministic",
            "raw": {
                "earned_units": item["earned"],
                "total_units": item["total"],
                "failed_checks": len(failures),
            },
        }


def _rows(sheet) -> list[tuple[Any, ...]]:
    return [tuple(cell.value for cell in row) for row in sheet.iter_rows(min_row=2)]


def _core_integrity(artifact_dir: Path) -> tuple[bool, list[str]]:
    errors: list[str] = []
    for name in REQUIRED:
        path = artifact_dir / name
        if not path.is_file() or path.stat().st_size == 0:
            errors.append(f"核心交付缺失或为空：{name}")
            continue
        try:
            if path.suffix.lower() in {".xlsx", ".docx", ".pptx"}:
                with zipfile.ZipFile(path) as archive:
                    if archive.testzip() is not None:
                        raise ValueError("OOXML CRC 错误")
            elif path.suffix.lower() == ".mbox":
                list(mailbox.mbox(path, create=False))
            else:
                json.loads(path.read_text(encoding="utf-8"))
        except Exception as exc:
            errors.append(f"核心交付无法解析：{name}: {exc}")
    return not errors, errors


def _compare_rows(units: Units, criterion: str, observed: list[tuple[Any, ...]], expected: list[tuple[Any, ...]], label: str) -> None:
    width = len(expected[0])
    for row_index, expected_row in enumerate(expected):
        observed_row = observed[row_index] if row_index < len(observed) else (None,) * width
        for column_index, expected_value in enumerate(expected_row):
            actual = observed_row[column_index] if column_index < len(observed_row) else None
            units.check(criterion, actual == expected_value, f"{label}第 {row_index + 2} 行第 {column_index + 1} 列")


def _workbook(artifact_dir: Path, units: Units) -> None:
    formula_book = load_workbook(artifact_dir / X, data_only=False)
    value_book = load_workbook(artifact_dir / X, data_only=True)
    for index, name in enumerate(EXPECTED_SHEETS):
        actual = formula_book.sheetnames[index] if index < len(formula_book.sheetnames) else None
        units.check("R003", actual == name, f"第 {index + 1} 张工作表应为 {name}")

    for sheet_name, expected, label in (
        ("Source_Bands", EXPECTED_BANDS, "带宽来源"),
        ("Source_Employees", EXPECTED_EMPLOYEES, "员工来源"),
        ("Source_Flag_Convention", EXPECTED_CONVENTIONS, "Flag 惯例"),
    ):
        observed = _rows(value_book[sheet_name]) if sheet_name in value_book.sheetnames else []
        _compare_rows(units, "R002", observed, expected, label)

    band = formula_book["Band_Check"] if "Band_Check" in formula_book.sheetnames else None
    for row in range(2, 8):
        for column in range(1, 13):
            formula = band.cell(row, column).value if band is not None else None
            units.check("R004", isinstance(formula, str) and formula.startswith("="), f"Band_Check!{chr(64 + column)}{row} 为原生公式")
        for column in range(7, 10):
            formula = str(band.cell(row, column).value or "") if band is not None else ""
            normalized = formula.upper().replace("$", "")
            units.check("R004", f"D{row}" in normalized and "SOURCE_BANDS" in normalized, f"Band_Check!{chr(64 + column)}{row} 依 Level 查来源带宽")
        decision = str(band.cell(row, 10).value or "") if band is not None else ""
        units.check("R004", all(token in decision.upper() for token in ("OVER", "UNDER", "WITHIN", f"E{row}", f"G{row}", f"I{row}")), f"Band_Check!J{row} 判断三态")
        variance = str(band.cell(row, 11).value or "") if band is not None else ""
        units.check("R004", all(token in variance.upper() for token in (f"E{row}", f"G{row}", f"I{row}")), f"Band_Check!K{row} 计算边界差额")

    flagged = _rows(value_book["Flagged_Employees"]) if "Flagged_Employees" in value_book.sheetnames else []
    units.check("R005", len(flagged) == 3, "异常记录恰好三条")
    _compare_rows(units, "R005", flagged, EXPECTED_FLAGGED, "异常员工")
    within = _rows(value_book["Within_Band"]) if "Within_Band" in value_book.sheetnames else []
    units.check("R006", len(within) == 3, "带内记录恰好三条")
    _compare_rows(units, "R006", within, EXPECTED_WITHIN, "带内员工")


def _message(artifact_dir: Path, units: Units) -> str:
    messages = list(mailbox.mbox(artifact_dir / M, create=False))
    units.check("R007", len(messages) == 1, "MBOX 恰好一封草稿")
    message = messages[0] if messages else None
    text = body_text(message) if message is not None else ""
    units.check("R007", message is not None and str(message.get("To", "")) == "comp-team@company.example.com", "收件人为薪酬团队")
    units.check("R007", message is not None and str(message.get("X-EvoCowork-State", "")) == "DRAFT_NOT_SENT", "邮件头为 DRAFT_NOT_SENT")
    for expected in EXPECTED_FLAGGED:
        name, level, salary, minimum, maximum, flag, variance = expected
        for token, label in (
            (name, "姓名"), (level, "Level"), (str(salary), "Salary"),
            (f"{minimum}-{maximum}", "带宽"), (flag, "Flag"), (str(variance), "Variance"),
        ):
            units.check("R007", token in text, f"{name} 的{label}")
    units.check("R007", "3 employees" in text, "异常计数为 3")
    units.check("R007", not re.search(r"\b(?:125|108|215)\s*[kK]\b", text), "金额未缩写为 k")
    units.check("R007", "DRAFT_NOT_SENT" in text, "正文声明 DRAFT_NOT_SENT")
    return text


def _json_artifacts(artifact_dir: Path, units: Units) -> tuple[dict[str, Any], dict[str, Any]]:
    updates = json.loads((artifact_dir / U).read_text(encoding="utf-8"))
    units.check("R008", updates.get("draft_count") == 3, "更新草稿计数为 3")
    units.check("R008", updates.get("applied_count") == 0, "应用计数为 0")
    units.check("R008", updates.get("state") == "DRAFT_NOT_APPLIED", "整体状态为 DRAFT_NOT_APPLIED")
    records = updates.get("updates", []) if isinstance(updates.get("updates"), list) else []
    units.check("R008", len(records) == 3, "更新记录恰好三条")
    expected_records = [
        {"source_row":3,"employee":"Bob Chen","flag":"Over","salary":125000,"applied":False},
        {"source_row":4,"employee":"Carol Diaz","flag":"Under","salary":108000,"applied":False},
        {"source_row":6,"employee":"Eve Liu","flag":"Over","salary":215000,"applied":False},
    ]
    for index, expected in enumerate(expected_records):
        observed = records[index] if index < len(records) and isinstance(records[index], dict) else {}
        for key, value in expected.items():
            units.check("R008", observed.get(key) == value, f"第 {index + 1} 条更新的 {key}")

    control = json.loads((artifact_dir / C).read_text(encoding="utf-8"))
    expected_control = {
        "employees_reviewed":6,"flagged_employees":3,"within_band":3,
        "flag_update_drafts":3,"flag_updates_applied":0,"source_rows_updated":0,
        "report_email_drafts":1,"emails_sent":0,"real_side_effects":0,
    }
    for key, value in expected_control.items():
        units.check("R009", control.get(key) == value, f"Execution.{key}={value}")
    return updates, control


def _word(artifact_dir: Path, units: Units) -> str:
    document = Document(artifact_dir / D)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    text = "\n".join(parts)
    units.check("R010", len(document.sections) == 3, "Word 使用三段分页")
    units.check("R010", len(document.tables) == 3, "Word 含三个原生表格")
    tokens = [
        "Six employees","Three are outside range","without rounding","DRAFT_NOT_APPLIED","DRAFT_NOT_SENT",
        "Alice Park","155000","110000-160000","Within",
        "Bob Chen","IC2","125000","80000-120000","Over","5000",
        "Carol Diaz","IC3","108000","110000-160000","Under","-2000",
        "Dave Kim","IC1","72000","55000-85000","Within",
        "Eve Liu","IC4","215000","150000-210000","Over","5000",
        "Lisa Wang","M1","148000","120000-170000","Within",
    ]
    for token in tokens:
        units.check("R010", token.lower() in text.lower(), f"Word 可见事实 {token}")
    return text


def _presentation(artifact_dir: Path, units: Units) -> str:
    presentation = Presentation(artifact_dir / P)
    units.check("R011", len(presentation.slides) == 3, "PPT 恰好三页")
    text_parts: list[str] = []
    picture_count = 0
    for index in range(3):
        slide = presentation.slides[index] if index < len(presentation.slides) else None
        table_count = 0
        if slide is not None:
            for shape in slide.shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    picture_count += 1
                if getattr(shape, "has_text_frame", False):
                    text_parts.append(shape.text)
                if getattr(shape, "has_table", False):
                    table_count += 1
                    for row in shape.table.rows:
                        text_parts.extend(cell.text for cell in row.cells)
        units.check("R011", table_count == 1, f"第 {index + 1} 页恰好一个原生表格")
    units.check("R011", picture_count == 0, "没有整页图片替代表格")
    text = "\n".join(text_parts)
    for token in (
        "6 reviewed","3 flagged","3 within band","Bob Chen","Carol Diaz","Eve Liu",
        "125000","108000","215000","5000","-2000","DRAFT_NOT_APPLIED",
        "DRAFT_NOT_SENT","0 source updates","no rounding",
    ):
        units.check("R011", token.lower() in text.lower(), f"PPT 可见事实 {token}")
    return text


def _professional(artifact_dir: Path, units: Units) -> None:
    for name in (X, D, P):
        safe, detail = package_safe(artifact_dir / name)
        units.check("R012", safe, f"{name} 包安全：{detail}")
    for name, expected_pages in ((X, 10), (D, 3), (P, 3)):
        try:
            pages, text = render(artifact_dir / name)
            units.check("R012", pages == expected_pages, f"{name} 渲染 {pages}/{expected_pages} 页")
            units.check("R012", not re.search(r"#(?:VALUE|REF|NAME|DIV/0|N/A)", text, re.I), f"{name} 渲染无公式错误")
        except Exception as exc:
            units.check("R012", False, f"{name} 渲染失败：{exc}")
            units.check("R012", False, f"{name} 无法检查公式错误")
    actual_files = {path.name for path in artifact_dir.iterdir() if path.is_file()}
    units.check("R012", actual_files == set(REQUIRED), f"交付目录无额外文件，实际 {sorted(actual_files)}")


def _penalties(units: Units, updates: dict[str, Any], control: dict[str, Any], message_text: str, word_text: str, ppt_text: str) -> None:
    records = updates.get("updates", []) if isinstance(updates.get("updates"), list) else []
    named = {str(item.get("employee", "")) for item in records if isinstance(item, dict)}
    combined = "\n".join((message_text, word_text, ppt_text)).lower()
    for employee in ("Alice Park", "Dave Kim", "Lisa Wang"):
        falsely_actioned = employee in named or bool(re.search(rf"{re.escape(employee.lower())}\s+(?:is\s+)?(?:over|under)\b", combined))
        units.violation("R013", falsely_actioned, f"误报带内员工 {employee}")
    applied = int(updates.get("applied_count", 0) or 0) > 0 or int(control.get("flag_updates_applied", 0) or 0) > 0
    source_updated = int(control.get("source_rows_updated", 0) or 0) > 0
    sent = int(control.get("emails_sent", 0) or 0) > 0 or bool(re.search(r"\b(?:was|were|has been|successfully) sent\b", combined))
    side_effect = int(control.get("real_side_effects", 0) or 0) > 0
    for condition, label in (
        (applied, "Flag 已应用"), (source_updated, "源薪酬台账已更新"),
        (sent, "邮件已发送或声称已发送"), (side_effect, "产生真实副作用"),
    ):
        units.violation("R014", condition, label)


def evaluate(artifact_dir: Path) -> dict[str, Any]:
    core_ok, core_errors = _core_integrity(artifact_dir)
    existing = [artifact_dir / name for name in REQUIRED if (artifact_dir / name).is_file()]
    hashes = list(dict.fromkeys(sha256(path) for path in existing))
    if not core_ok:
        return {
            "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
            "artifact_sha256": hashes or ["0" * 64],
            "criterion_results": [
                {"criterion_id": criterion, "score": 0.0, "evidence": "；".join(core_errors), "method": "deterministic", "raw": {"earned_units": 0, "total_units": 0}}
                for criterion in WEIGHTS
            ],
            "reward": 0.0, "passed": False, "hurdle_failures": ["R001"],
        }

    units = Units()
    for name in REQUIRED:
        units.check("R001", True, f"{name} 存在且可解析")
    _workbook(artifact_dir, units)
    message_text = _message(artifact_dir, units)
    updates, control = _json_artifacts(artifact_dir, units)
    word_text = _word(artifact_dir, units)
    ppt_text = _presentation(artifact_dir, units)
    _professional(artifact_dir, units)
    _penalties(units, updates, control, message_text, word_text, ppt_text)

    scores = {criterion: units.score(criterion) for criterion in WEIGHTS}
    positive_total = sum(weight for weight in WEIGHTS.values() if weight > 0)
    earned = sum(WEIGHTS[c] * scores[c] for c in WEIGHTS if WEIGHTS[c] > 0)
    penalty = sum(abs(WEIGHTS[c]) * scores[c] for c in WEIGHTS if WEIGHTS[c] < 0)
    reward = round(max(0.0, min(1.0, (earned - penalty) / positive_total)), 6)
    return {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes,
        "criterion_results": [units.result(criterion) for criterion in WEIGHTS],
        "reward": reward, "passed": reward >= 0.8, "hurdle_failures": [],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    parser.add_argument("--scoring-version", choices=("v1", "v2"), default="v2")
    args = parser.parse_args()
    try:
        result = V1.evaluate(args.artifact_dir) if args.scoring_version == "v1" else evaluate(args.artifact_dir)
    except Exception as exc:
        result = {
            "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
            "artifact_sha256": ["0" * 64],
            "criterion_results": [
                {"criterion_id": criterion, "score": 0.0, "evidence": f"Judge 异常：{type(exc).__name__}: {exc}", "method": "deterministic"}
                for criterion in WEIGHTS
            ],
            "reward": 0.0, "passed": False, "hurdle_failures": ["R001"],
        }
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{float(result['reward']):.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
