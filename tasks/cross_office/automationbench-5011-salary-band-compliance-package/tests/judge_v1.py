#!/usr/bin/env python3
"""AutomationBench 5011 薪酬带区间与精确金额的 v1 严格 Judge。"""
import argparse,json,mailbox,re
from pathlib import Path
from openpyxl import load_workbook
from office_checks import body_text,docx_text,package_safe,pptx_text,render,sha256
TASK="cowork-evo/automationbench-5011-salary-band-compliance-package";HASHES={"Salary_Band_Audit_Source.xlsx":"c957d578ace7081a3bd51eec0465b3886760a1405b382c9db0c09fbf47d8efce","Workflow_Context.json":"9ea425af08fa2ebb0a5d8cc2b5dfcaeaf3b9b8f35b06a184eda50466be34a60d"};X="Salary_Band_Compliance_Audit.xlsx";D="Salary_Band_Compliance_Brief.docx";P="Salary_Band_Compliance_Review.pptx";M="Compensation_Team_Report_Draft.mbox";U="Salary_Flag_Update_Drafts.json";C="Salary_Band_Execution_Control.json";REQUIRED={X,D,P,M,U,C};WEIGHTS={f"R{i:03d}":v for i,v in enumerate([6,8,10,10,10,8,9,8,8,8,7,8],1)};CAPS={k:(0 if k=="R001" else .25) for k in WEIGHTS if k!="R012"}
class Checks:
 def __init__(self):self.v={k:{"passed":True,"details":[]} for k in WEIGHTS}
 def require(self,k,c,d):
  if not c:self.v[k]["passed"]=False;self.v[k]["details"].append(d) if d not in self.v[k]["details"] else None
 def ok(self,k):return self.v[k]["passed"]
def root():return Path(__file__).resolve().parents[1]
def inp():return Path("/app/input_files") if Path("/app/input_files").is_dir() else root()/"data/input_files"
def rows(ws):return [tuple(c.value for c in r) for r in ws.iter_rows(min_row=2)]
def source(q):
 for n,h in HASHES.items():q.require("R001",(inp()/n).is_file() and sha256(inp()/n)==h,f"输入{n}")
 wb=load_workbook(inp()/"Salary_Band_Audit_Source.xlsx",read_only=True,data_only=True);b=rows(wb["Salary Bands"]);e=rows(wb["Employee Compensation"]);c=rows(wb["Audit Flag Convention"]);q.require("R002",[r[1] for r in b]==["IC1","IC2","IC3","IC4","M1"] and [(r[3],r[4],r[5]) for r in b]==[(55000,70000,85000),(80000,100000,120000),(110000,135000,160000),(150000,180000,210000),(120000,145000,170000)],"5档带宽")
 q.require("R002",[(r[1],r[3],r[4]) for r in e]==[("Alice Park","IC3",155000),("Bob Chen","IC2",125000),("Carol Diaz","IC3",108000),("Dave Kim","IC1",72000),("Eve Liu","IC4",215000),("Lisa Wang","M1",148000)] and c==[(2,"Salary above band maximum","Over"),(3,"Salary below band minimum","Under")],"6人薪酬与标记惯例")
def workbook(p,q):
 wb=load_workbook(p,data_only=False);q.require("R003",wb.sheetnames==["Source_Bands","Source_Employees","Source_Flag_Convention","Evidence_Register","Band_Check","Flagged_Employees","Within_Band","Flag_Update_Drafts","Report_Draft","Control_Summary"],"10表");f=[c.value for r in wb["Band_Check"].iter_rows(min_row=2) for c in r if isinstance(c.value,str) and c.value.startswith("=")];q.require("R004",len(f)==72,f"公式{len(f)}");j="\n".join(f)
 for t in ("VLOOKUP","Source_Bands!$B$2:$F$6",'"Over"','"Under"','"Within"','E2-I2','E2-G2','Source_Employees!F7'):q.require("R004",t in j,f"公式缺{t}")
 flagged=rows(wb["Flagged_Employees"]);q.require("R005",flagged==[("Bob Chen","IC2",125000,80000,120000,"Over",5000),("Carol Diaz","IC3",108000,110000,160000,"Under",-2000),("Eve Liu","IC4",215000,150000,210000,"Over",5000)],"3条异常精确金额")
 within=rows(wb["Within_Band"]);q.require("R006",[r[0] for r in within]==["Alice Park","Dave Kim","Lisa Wang"] and all(r[5]=="Within" for r in within),"3人带内")
def artifacts(a,q):
 ms=list(mailbox.mbox(a/M,create=False));text=body_text(ms[0]) if len(ms)==1 else "";q.require("R007",len(ms)==1 and str(ms[0].get("To"))=="comp-team@company.example.com" and all(t in text for t in ("3 employees","Bob Chen","125000","80000-120000","Carol Diaz","108000","110000-160000","Eve Liu","215000","150000-210000","DRAFT_NOT_SENT")),"邮件精确金额")
 u=json.loads((a/U).read_text());q.require("R008",u.get("draft_count")==3 and u.get("applied_count")==0 and [(x["employee"],x["flag"],x["salary"],x["applied"]) for x in u.get("updates",[])]==[("Bob Chen","Over",125000,False),("Carol Diaz","Under",108000,False),("Eve Liu","Over",215000,False)],"标记更新草稿")
 c=json.loads((a/C).read_text());exp={"employees_reviewed":6,"flagged_employees":3,"within_band":3,"flag_update_drafts":3,"flag_updates_applied":0,"source_rows_updated":0,"report_email_drafts":1,"emails_sent":0,"real_side_effects":0};q.require("R009",all(c.get(k)==v for k,v in exp.items()),"执行边界")
def office(a,q):
 text,s,t=docx_text(a/D);q.require("R010",s==3 and t==3,"Word结构")
 for x in ("Six employees","Bob Chen","125000","Carol Diaz","108000","Eve Liu","215000","without rounding","DRAFT_NOT_APPLIED"):q.require("R010",x.lower() in text.lower(),f"Word缺{x}")
 q.require("R010","alice park is over" not in text.lower(),"Word不得误报Alice")
 text,s,t,pic=pptx_text(a/P);q.require("R011",s==3 and t==[1,1,1] and pic==0,"PPT结构")
 for x in ("3 flagged","Bob Chen","Carol Diaz","Eve Liu","125000","108000","215000","DRAFT_NOT_APPLIED"):q.require("R011",x in text,f"PPT缺{x}")
 for n in (X,D,P):safe,d=package_safe(a/n);q.require("R012",safe,f"安全{n}:{d}")
def finish(q,pages):
 h=[k for k in CAPS if not q.ok(k)];base=sum(v for k,v in WEIGHTS.items() if q.ok(k))/100;cap=min([1]+[CAPS[k] for k in h]);pen=(.3 if any(not q.ok(k) for k in ("R002","R003","R004","R005","R006")) else 0)+(.3 if any(not q.ok(k) for k in ("R007","R008","R009")) else 0)+(.25 if any(not q.ok(k) for k in ("R001","R010","R011","R012")) else 0);reward=round(max(0,min(base,cap)-pen),6);return {"task_id":TASK,"reward":reward,"passed":reward>=.8 and not h,"criteria":q.v,"hurdle_failures":h,"source_truth":{"reviewed":6,"flagged":3,"within":3,"source_updates":0,"real_side_effects":0},"render_pages":pages}
def evaluate(a):
 q=Checks();source(q);files={p.name for p in a.iterdir() if p.is_file()} if a.is_dir() else set();q.require("R001",files==REQUIRED,f"集合{files}")
 if files!=REQUIRED:return finish(q,{})
 for fn,ks in ((lambda:workbook(a/X,q),("R003","R004","R005","R006","R012")),(lambda:artifacts(a,q),("R007","R008","R009")),(lambda:office(a,q),("R010","R011","R012"))):
  try:fn()
  except Exception as e:
   for k in ks:q.require(k,False,str(e))
 pages={}
 for n,e in ((X,10),(D,3),(P,3)):
  try:c,t=render(a/n);pages[n]=c;q.require("R012",c==e,f"页数{n}:{c}");q.require("R012",not re.search(r"#(?:VALUE|REF|NAME|DIV/0|N/A)",t,re.I),f"公式错{n}")
  except Exception as x:q.require("R012",False,str(x))
 return finish(q,pages)
def main():
 p=argparse.ArgumentParser();p.add_argument("--artifact-dir",type=Path,required=True);p.add_argument("--result",type=Path,required=True);p.add_argument("--reward-file",type=Path,required=True);a=p.parse_args();r=evaluate(a.artifact_dir);a.result.parent.mkdir(parents=True,exist_ok=True);a.reward_file.parent.mkdir(parents=True,exist_ok=True);a.result.write_text(json.dumps(r,ensure_ascii=False,indent=2)+"\n");a.reward_file.write_text(f"{r['reward']:.6f}\n");print(json.dumps(r,ensure_ascii=False,indent=2))
if __name__=="__main__":main()
