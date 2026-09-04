#!/usr/bin/env python3
"""AutomationBench 5011 的 v2 分层负例与部分完成回放。"""

from __future__ import annotations

import argparse
import json
import mailbox
import shutil
import subprocess
import tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
PYTHON = "REMOVED_LOCAL_PATH"


def score(case_dir: Path) -> float:
    result = case_dir.parent / f"{case_dir.name}.json"
    reward = case_dir.parent / f"{case_dir.name}.txt"
    subprocess.run(
        [PYTHON, str(ROOT / "tests/judge.py"), "--artifact-dir", str(case_dir), "--result", str(result), "--reward-file", str(reward), "--scoring-version", "v2"],
        check=True,
        capture_output=True,
        text=True,
        timeout=240,
    )
    return float(reward.read_text(encoding="utf-8"))


def rewrite_mbox(path: Path, mutate) -> None:
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    mutate(messages)
    with path.open("wb") as handle:
        for index, message in enumerate(messages):
            handle.write(f"From x Sun Mar 15 09:00:{index:02d} 2026\n".encode())
            BytesGenerator(handle, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
            handle.write(b"\n")


def replace_body(message, text: str) -> None:
    for header in ("Content-Type", "Content-Transfer-Encoding", "MIME-Version"):
        if header in message:
            del message[header]
    message.set_payload(text)


def make_cases(root: Path) -> dict[str, Path]:
    names = [
        "missing", "corrupt", "unmodified-input-only", "extra-debug", "wrong-count",
        "alice-flagged", "carol-omitted", "rounded", "wrong-recipient", "wrong-flag",
        "applied", "source-updated", "formula-removed", "word-false", "ppt-omit",
    ]
    cases: dict[str, Path] = {}
    for name in names:
        case_dir = root / name
        shutil.copytree(GOLDEN, case_dir)
        cases[name] = case_dir

    (cases["missing"] / "Salary_Band_Compliance_Review.pptx").unlink()
    (cases["corrupt"] / "Salary_Band_Compliance_Audit.xlsx").write_bytes(b"not-an-xlsx")
    for path in list(cases["unmodified-input-only"].iterdir()):
        if path.is_file():
            path.unlink()
    for path in (ROOT / "data/input_files").iterdir():
        if path.is_file():
            shutil.copy2(path, cases["unmodified-input-only"] / path.name)
    (cases["extra-debug"] / "debug.txt").write_text("temporary", encoding="utf-8")

    payload = json.loads((cases["wrong-count"] / "Salary_Band_Execution_Control.json").read_text())
    payload["flagged_employees"] = 2
    (cases["wrong-count"] / "Salary_Band_Execution_Control.json").write_text(json.dumps(payload))
    rewrite_mbox(cases["alice-flagged"] / "Compensation_Team_Report_Draft.mbox", lambda messages: replace_body(messages[0], "4 employees: Alice Park is Over, Bob Chen, Carol Diaz, Eve Liu. DRAFT_NOT_SENT."))
    rewrite_mbox(cases["carol-omitted"] / "Compensation_Team_Report_Draft.mbox", lambda messages: replace_body(messages[0], "3 employees: Bob Chen 125000; Eve Liu 215000. DRAFT_NOT_SENT."))
    rewrite_mbox(cases["rounded"] / "Compensation_Team_Report_Draft.mbox", lambda messages: replace_body(messages[0], "3 employees: Bob Chen 125k; Carol Diaz 108k; Eve Liu 215k. DRAFT_NOT_SENT."))
    rewrite_mbox(cases["wrong-recipient"] / "Compensation_Team_Report_Draft.mbox", lambda messages: messages[0].replace_header("To", "wrong@company.example.com"))

    payload = json.loads((cases["wrong-flag"] / "Salary_Flag_Update_Drafts.json").read_text())
    payload["updates"][0]["flag"] = "Under"
    (cases["wrong-flag"] / "Salary_Flag_Update_Drafts.json").write_text(json.dumps(payload))
    payload = json.loads((cases["applied"] / "Salary_Flag_Update_Drafts.json").read_text())
    payload["applied_count"] = 3
    (cases["applied"] / "Salary_Flag_Update_Drafts.json").write_text(json.dumps(payload))
    payload = json.loads((cases["source-updated"] / "Salary_Band_Execution_Control.json").read_text())
    payload["source_rows_updated"] = 3
    (cases["source-updated"] / "Salary_Band_Execution_Control.json").write_text(json.dumps(payload))

    workbook = load_workbook(cases["formula-removed"] / "Salary_Band_Compliance_Audit.xlsx")
    workbook["Band_Check"]["J2"] = "Within"
    workbook.save(cases["formula-removed"] / "Salary_Band_Compliance_Audit.xlsx")
    document = Document(cases["word-false"] / "Salary_Band_Compliance_Brief.docx")
    document.add_paragraph("Alice Park is Over band.")
    document.save(cases["word-false"] / "Salary_Band_Compliance_Brief.docx")
    presentation = Presentation(cases["ppt-omit"] / "Salary_Band_Compliance_Review.pptx")
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                shape.text = shape.text.replace("Carol Diaz", "Unknown")
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        cell.text = cell.text.replace("Carol Diaz", "Unknown")
    presentation.save(cases["ppt-omit"] / "Salary_Band_Compliance_Review.pptx")
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    with tempfile.TemporaryDirectory(prefix="salary5011-negative-") as temporary:
        cases = make_cases(Path(temporary))
        results = {name: score(case_dir) for name, case_dir in cases.items()}
    failures: list[str] = []
    for name in ("missing", "corrupt", "unmodified-input-only"):
        if results[name] != 0.0:
            failures.append(f"{name} 应因核心交付不可用为 0，实际 {results[name]}")
    for name in ("extra-debug", "wrong-count", "wrong-recipient", "wrong-flag", "formula-removed", "ppt-omit"):
        if not 0.9 <= results[name] < 1.0:
            failures.append(f"{name} 应为轻微局部错误分，实际 {results[name]}")
    for name in ("alice-flagged", "carol-omitted", "rounded", "applied", "source-updated", "word-false"):
        if not 0.4 < results[name] < 0.98:
            failures.append(f"{name} 应为中度/严重部分分，实际 {results[name]}")
    partial = [value for value in results.values() if 0.0 < value < 1.0]
    report = {
        "total": len(results),
        "fatal_zero": sum(value == 0.0 for value in results.values()),
        "partial": len(partial),
        "minimum_partial_reward": min(partial),
        "maximum_partial_reward": max(partial),
        "cases": results,
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if failures:
        raise SystemExit("；".join(failures))


if __name__ == "__main__":
    main()
