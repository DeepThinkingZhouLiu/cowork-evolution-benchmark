#!/usr/bin/env python3
"""只静态化 AutomationBench 5057；不是通用造题器。"""
import hashlib,importlib.util,json,shutil,sys,types
from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.worksheet.table import Table,TableStyleInfo
ROOT=Path(__file__).resolve().parents[1];UP=Path(__file__).with_name("upstream-hr-tasks.py");SNAP=Path(__file__).with_name("upstream-task-5057.json");DATA=ROOT/"data/input_files";ENV=ROOT/"environment/input_files";UPSTREAM_SHA="c173a36a03f43096bd122577584aa32f9df379d8760aa46c43be04c85d101779"
def load_task():
 if hashlib.sha256(UP.read_bytes()).hexdigest()!=UPSTREAM_SHA:raise ValueError("上游哈希错误")
 d=types.ModuleType("datasets");d.Dataset=object;sys.modules["datasets"]=d
 for n in ("automationbench","automationbench.domains","automationbench.domains.hr"):sys.modules[n]=types.ModuleType(n)
 noise=types.ModuleType("automationbench.domains.hr._noise");noise.apply_noise=lambda x:x;sys.modules[noise.__name__]=noise
 s=importlib.util.spec_from_file_location("automationbench_5057",UP);m=importlib.util.module_from_spec(s);s.loader.exec_module(m);t=m.get_hr_multi_app_new_hire_setup_task();assert t["example_id"]==5057 and t["task"]=="hr.multi_app_new_hire_setup";return t
def audit(t):
 i=t["info"];s=i["initial_state"];assert s["meta"]["current_time"]=="2026-03-22T09:00:00Z" and len(i["assertions"])==11
 r=s["google_sheets"]["spreadsheets"][0]["worksheets"][0]["rows"];assert len(r)==2 and [x["cells"]["Name"] for x in r]==["Alicia Fernandez","Tyrone Washington"] and [x["cells"]["Start Date"] for x in r]==["2026-03-24","2026-03-25"]
 assert r[0]["cells"]["Manager"]=="David Okonkwo" and r[1]["cells"]["Manager"]=="Lisa Wang" and s["google_drive"]["folders"]==[{"id":"folder_employees","name":"Employee Documents"}]
def write_xlsx(rows,p):
 wb=Workbook();ws=wb.active;ws.title="This Week";h=["Source Row","Name","Email","Department","Manager","Start Date","Location","Role"];ws.append(h)
 for r in rows:ws.append([r["row_id"],r["cells"]["Name"],r["cells"]["Email"],r["cells"]["Department"],r["cells"]["Manager"],date.fromisoformat(r["cells"]["Start Date"]),r["cells"]["Location"],r["cells"]["Role"]])
 for i,w in enumerate([10,26,42,22,24,18,18,30],1):ws.column_dimensions[ws.cell(1,i).column_letter].width=w
 ws.freeze_panes="A2";ws.sheet_view.showGridLines=False
 for c in ws[1]:c.fill=PatternFill("solid",fgColor="17365D");c.font=Font(name="Carlito",bold=True,color="FFFFFF");c.alignment=Alignment(horizontal="center",wrap_text=True)
 for row in ws.iter_rows(min_row=2):
  for c in row:c.font=Font(name="Carlito",size=9);c.alignment=Alignment(vertical="top",wrap_text=True)
  row[5].number_format="yyyy-mm-dd"
 t=Table(displayName="NewHireChecklistSource",ref=ws.dimensions);t.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True);ws.add_table(t);ws.page_setup.orientation="landscape";ws.page_setup.fitToWidth=1;ws.sheet_properties.pageSetUpPr.fitToPage=True;wb.save(p)
def main():
 t=load_task();audit(t);s=t["info"]["initial_state"];DATA.mkdir(parents=True,exist_ok=True);ENV.mkdir(parents=True,exist_ok=True);snap=json.dumps(t,ensure_ascii=False,indent=2)+"\n"
 if SNAP.exists() and SNAP.read_text()!=snap:raise ValueError("单题快照漂移")
 SNAP.write_text(snap);write_xlsx(s["google_sheets"]["spreadsheets"][0]["worksheets"][0]["rows"],DATA/"Full_Onboarding_Checklist.xlsx");(DATA/"Initial_Application_State.json").write_text(json.dumps({k:s[k] for k in ("bamboohr","google_drive","jira")},ensure_ascii=False,indent=2)+"\n");(DATA/"Workflow_Context.json").write_text(json.dumps({"meta":s["meta"]},indent=2)+"\n")
 for f in DATA.iterdir():
  if f.is_file():shutil.copy2(f,ENV/f.name)
if __name__=="__main__":main()
