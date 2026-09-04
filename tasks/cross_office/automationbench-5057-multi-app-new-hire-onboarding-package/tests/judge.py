#!/usr/bin/env python3
"""AutomationBench 5057 的 v2 原子部分分 Judge；v1 保存在 judge_v1.py。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import mailbox
import re
import sys
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from office_checks import body_text, package_safe, render


TASK_ID = "cowork-evo/automationbench-5057-multi-app-new-hire-onboarding-package"
RUBRIC_VERSION = "automationbench-5057-v0.2"
JUDGE_VERSION = "automationbench-5057-atomic-partial-v2.0"
X = "New_Hire_Onboarding_Audit.xlsx"
D = "New_Hire_Onboarding_Brief.docx"
P = "New_Hire_Onboarding_Review.pptx"
B = "BambooHR_Profile_Create_Drafts.json"
F = "Drive_Folder_Create_Drafts.json"
J = "Jira_Onboarding_Ticket_Drafts.json"
M = "New_Hire_Welcome_Drafts.mbox"
C = "Onboarding_Execution_Control.json"
REQUIRED = (X, D, P, B, F, J, M, C)
WEIGHTS = {
    "R001": 6, "R002": 8, "R003": 10, "R004": 10, "R005": 10,
    "R006": 8, "R007": 9, "R008": 8, "R009": 8, "R010": 8,
    "R011": 7, "R012": 8, "R013": -15, "R014": -15,
}
EXPECTED_SHEETS = [
    "Source_Checklist", "Source_System_State", "Evidence_Register", "Hire_Decisions",
    "BambooHR_Drafts", "Drive_Drafts", "Jira_Drafts", "Welcome_Drafts",
    "Manager_Exclusion", "Cross_System_Reconciliation", "Readiness_Matrix", "Control_Summary",
]
HIRES = [
    {
        "row": 2, "name": "Alicia Fernandez", "email": "alicia.fernandez@company.example.com",
        "department": "Engineering", "manager": "David Okonkwo", "start": "2026-03-24",
        "location": "Austin", "role": "Frontend Developer",
    },
    {
        "row": 3, "name": "Tyrone Washington", "email": "tyrone.washington@company.example.com",
        "department": "Sales", "manager": "Lisa Wang", "start": "2026-03-25",
        "location": "Remote", "role": "Account Executive",
    },
]


def _load_v1():
    path = Path(__file__).with_name("judge_v1.py")
    spec = importlib.util.spec_from_file_location("automationbench_5057_judge_v1", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v1 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


V1 = _load_v1()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


class Units:
    def __init__(self) -> None:
        self.values: dict[str, dict[str, Any]] = {
            criterion: {"earned": 0.0, "total": 0.0, "failed": []}
            for criterion in WEIGHTS
        }

    def check(self, criterion: str, condition: bool, label: str, count: float = 1.0) -> None:
        item = self.values[criterion]
        item["total"] += count
        if condition:
            item["earned"] += count
        else:
            item["failed"].append(label)

    def violation(self, criterion: str, present: bool, label: str) -> None:
        item = self.values[criterion]
        item["total"] += 1.0
        if present:
            item["earned"] += 1.0
            item["failed"].append(label)

    def score(self, criterion: str) -> float:
        item = self.values[criterion]
        return round(item["earned"] / item["total"], 6) if item["total"] else 0.0

    def result(self, criterion: str) -> dict[str, Any]:
        item = self.values[criterion]
        failures = item["failed"]
        if WEIGHTS[criterion] < 0:
            evidence = "发现违规：" + "；".join(failures) if failures else "未发现本项违规"
        else:
            evidence = "未通过：" + "；".join(failures) if failures else f"全部 {int(item['total'])} 个任务特定检查单元通过"
        return {
            "criterion_id": criterion,
            "score": self.score(criterion),
            "evidence": evidence,
            "method": "deterministic",
            "raw": {
                "earned_units": item["earned"],
                "total_units": item["total"],
                "failed_checks": len(failures),
            },
        }


def _core_integrity(artifact_dir: Path) -> tuple[bool, list[str]]:
    errors: list[str] = []
    for name in REQUIRED:
        path = artifact_dir / name
        if not path.is_file() or path.stat().st_size == 0:
            errors.append(f"核心交付缺失或为空：{name}")
            continue
        try:
            if name == X:
                load_workbook(path, read_only=True, data_only=False).close()
            elif name == D:
                Document(path)
            elif name == P:
                Presentation(path)
            elif name == M:
                list(mailbox.mbox(path, create=False))
            else:
                json.loads(path.read_text(encoding="utf-8"))
        except Exception as exc:
            errors.append(f"核心交付无法解析：{name}: {exc}")
    return not errors, errors


def _rows(sheet) -> list[tuple[Any, ...]]:
    return [tuple(cell.value for cell in row) for row in sheet.iter_rows(min_row=2)]


def _workbook(artifact_dir: Path, units: Units) -> None:
    formula_book = load_workbook(artifact_dir / X, data_only=False)
    value_book = load_workbook(artifact_dir / X, data_only=True)

    for index, name in enumerate(EXPECTED_SHEETS):
        actual = formula_book.sheetnames[index] if index < len(formula_book.sheetnames) else None
        units.check("R003", actual == name, f"第 {index + 1} 张工作表应为 {name}")

    source_rows = _rows(value_book["Source_Checklist"]) if "Source_Checklist" in value_book.sheetnames else []
    for index, hire in enumerate(HIRES):
        expected = (
            hire["row"], hire["name"], hire["email"], hire["department"], hire["manager"],
            hire["start"], hire["location"], hire["role"],
        )
        observed = source_rows[index] if index < len(source_rows) else (None,) * len(expected)
        for column, value in enumerate(expected):
            actual = observed[column] if column < len(observed) else None
            if column == 5 and hasattr(actual, "strftime"):
                actual = actual.strftime("%Y-%m-%d")
            units.check("R002", actual == value, f"新人 {hire['name']} 的第 {column + 1} 个来源字段")

    system_rows = _rows(value_book["Source_System_State"]) if "Source_System_State" in value_book.sheetnames else []
    expected_systems = [
        ("BambooHR", "No actions", "None"),
        ("Google Drive", "One parent folder", "folder_employees / Employee Documents"),
        ("Jira", "No actions", "None"),
    ]
    for row_index, expected in enumerate(expected_systems):
        observed = system_rows[row_index] if row_index < len(system_rows) else (None, None, None)
        columns = range(3) if row_index < 2 else range(2)
        for column in columns:
            units.check("R002", observed[column] == expected[column], f"系统初始状态 {expected[0]} 字段 {column + 1}")

    decision = formula_book["Hire_Decisions"] if "Hire_Decisions" in formula_book.sheetnames else None
    for row in range(2, 4):
        for column in range(1, 17):
            value = decision.cell(row, column).value if decision is not None else None
            units.check("R004", isinstance(value, str) and value.startswith("="), f"Hire_Decisions!{chr(64 + column)}{row} 为原生公式")
        values = [str(decision.cell(row, column).value or "") if decision is not None else "" for column in range(1, 17)]
        normalized = [value.upper().replace("$", "") for value in values]
        units.check("R004", all(f"SOURCE_CHECKLIST!{chr(64 + column)}{row}" in normalized[column - 1] for column in range(1, 9)), f"第 {row} 行八个身份字段链接源清单")
        units.check("R004", "PROFILE_DRAFT:" in normalized[8] and f"B{row}" in normalized[8], f"第 {row} 行 BambooHR 档案链接姓名")
        units.check("R004", "FOLDER_EMPLOYEES/" in normalized[9] and f"B{row}" in normalized[9], f"第 {row} 行 Drive 父目录与姓名")
        units.check("R004", "ONBOARDING: " in normalized[10] and f"B{row}" in normalized[10], f"第 {row} 行 Jira Summary 链接姓名")
        units.check("R004", normalized[11] == f"=C{row}", f"第 {row} 行欢迎邮件收件人链接新人")
        units.check(
            "R004",
            all(token in "\n".join(normalized[12:]) for token in ("YES", "DAVID.OKONKWO", "LISA.WANG", "READY_DRAFT_NOT_EXECUTED", f"B{row}", f"E{row}")),
            f"第 {row} 行包含对账、经理排除、就绪和追踪语义",
        )

    reconciliation = _rows(value_book["Cross_System_Reconciliation"]) if "Cross_System_Reconciliation" in value_book.sheetnames else []
    expected_reconciliation = [
        ("Alicia Fernandez", 1, 1, 1, 1, "YES"),
        ("Tyrone Washington", 1, 1, 1, 1, "YES"),
    ]
    for row_index, expected in enumerate(expected_reconciliation):
        observed = reconciliation[row_index] if row_index < len(reconciliation) else (None,) * 6
        for column, value in enumerate(expected):
            units.check("R005", observed[column] == value, f"{expected[0]} 跨系统对账字段 {column + 1}")

    managers = _rows(value_book["Manager_Exclusion"]) if "Manager_Exclusion" in value_book.sheetnames else []
    expected_managers = [
        ("David Okonkwo", 0, 0, "Manager field; not a new hire"),
        ("Lisa Wang", 0, 0, "Manager field; not a new hire"),
    ]
    for row_index, expected in enumerate(expected_managers):
        observed = managers[row_index] if row_index < len(managers) else (None,) * 4
        for column, value in enumerate(expected):
            units.check("R006", observed[column] == value, f"{expected[0]} 经理排除字段 {column + 1}")


def _json_records(artifact_dir: Path, units: Units) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any], dict[str, Any]]:
    bamboo = json.loads((artifact_dir / B).read_text(encoding="utf-8"))
    units.check("R007", bamboo.get("draft_count") == 2, "BambooHR 草稿计数为 2")
    units.check("R007", bamboo.get("created_count") == 0, "BambooHR 创建计数为 0")
    profiles = bamboo.get("profiles", []) if isinstance(bamboo.get("profiles"), list) else []
    units.check("R007", len(profiles) == 2, "BambooHR 档案草稿恰好两条")
    profile_fields = ("name", "email", "department", "manager", "start_date", "location", "role", "state")
    for index, hire in enumerate(HIRES):
        observed = profiles[index] if index < len(profiles) and isinstance(profiles[index], dict) else {}
        expected = {
            "name": hire["name"], "email": hire["email"], "department": hire["department"],
            "manager": hire["manager"], "start_date": hire["start"], "location": hire["location"],
            "role": hire["role"], "state": "DRAFT_NOT_CREATED",
        }
        for key in profile_fields:
            units.check("R007", observed.get(key) == expected[key], f"BambooHR {hire['name']} 的 {key}")

    drive = json.loads((artifact_dir / F).read_text(encoding="utf-8"))
    units.check("R007", drive.get("draft_count") == 2, "Drive 草稿计数为 2")
    units.check("R007", drive.get("created_count") == 0, "Drive 创建计数为 0")
    folders = drive.get("folders", []) if isinstance(drive.get("folders"), list) else []
    units.check("R007", len(folders) == 2, "Drive 文件夹草稿恰好两条")
    for index, hire in enumerate(HIRES):
        observed = folders[index] if index < len(folders) and isinstance(folders[index], dict) else {}
        expected = {
            "name": hire["name"], "parent_id": "folder_employees",
            "parent_name": "Employee Documents", "state": "DRAFT_NOT_CREATED",
        }
        for key, value in expected.items():
            units.check("R007", observed.get(key) == value, f"Drive {hire['name']} 的 {key}")

    jira = json.loads((artifact_dir / J).read_text(encoding="utf-8"))
    units.check("R008", jira.get("draft_count") == 2, "Jira 草稿计数为 2")
    units.check("R008", jira.get("created_count") == 0, "Jira 创建计数为 0")
    tickets = jira.get("tickets", []) if isinstance(jira.get("tickets"), list) else []
    units.check("R008", len(tickets) == 2, "Jira 草稿恰好两条")
    for index, hire in enumerate(HIRES):
        ticket = tickets[index] if index < len(tickets) and isinstance(tickets[index], dict) else {}
        units.check("R008", ticket.get("summary") == f"Onboarding: {hire['name']}", f"Jira {hire['name']} Summary")
        description = str(ticket.get("description", ""))
        for key in ("name", "role", "department", "manager", "start", "location"):
            units.check("R008", hire[key] in description, f"Jira {hire['name']} Description 含 {key}")
        units.check("R008", ticket.get("state") == "DRAFT_NOT_CREATED", f"Jira {hire['name']} 未创建状态")

    messages = list(mailbox.mbox(artifact_dir / M, create=False))
    units.check("R008", len(messages) == 2, "欢迎 MBOX 恰好两封")
    for index, hire in enumerate(HIRES):
        message = messages[index] if index < len(messages) else None
        text = body_text(message) if message is not None else ""
        units.check("R008", message is not None and str(message.get("To", "")) == hire["email"], f"{hire['name']} 欢迎邮件收件人")
        for key in ("name", "department", "role", "start", "location", "manager"):
            units.check("R008", hire[key] in text, f"{hire['name']} 欢迎邮件含 {key}")
        units.check("R008", message is not None and str(message.get("X-EvoCowork-State", "")) == "DRAFT_NOT_SENT", f"{hire['name']} 欢迎邮件未发送")

    control = json.loads((artifact_dir / C).read_text(encoding="utf-8"))
    expected_control = {
        "new_hires": 2,
        "bamboo_profile_drafts": 2,
        "bamboo_profiles_created": 0,
        "drive_folder_drafts": 2,
        "drive_folders_created": 0,
        "jira_ticket_drafts": 2,
        "jira_tickets_created": 0,
        "welcome_email_drafts": 2,
        "welcome_emails_sent": 0,
        "manager_profiles_created": 0,
        "manager_welcome_emails_sent": 0,
        "external_actions_executed": 0,
        "real_side_effects": 0,
    }
    for key, value in expected_control.items():
        units.check("R009", control.get(key) == value, f"Execution.{key}={value}")
    return bamboo, drive, jira, control


def _word(artifact_dir: Path, units: Units) -> str:
    document = Document(artifact_dir / D)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    text = "\n".join(parts)
    units.check("R010", len(document.sections) == 3, "Word 使用三段分页")
    units.check("R010", len(document.tables) == 3, "Word 含三个原生表格")
    tokens = [
        "Alicia Fernandez", "alicia.fernandez@company.example.com", "Engineering", "David Okonkwo",
        "2026-03-24", "Austin", "Frontend Developer", "Tyrone Washington",
        "tyrone.washington@company.example.com", "Sales", "Lisa Wang", "2026-03-25", "Remote",
        "Account Executive", "BambooHR", "Google Drive", "Jira", "Gmail", "DRAFT_NOT_CREATED",
        "DRAFT_NOT_SENT", "appear only as managers", "no welcome email", "no BambooHR new-hire profile",
        "one BambooHR profile draft", "one Employee Documents folder draft", "one Jira onboarding ticket draft",
        "one welcome email draft", "offline only",
    ]
    for token in tokens:
        units.check("R010", token.lower() in text.lower(), f"Word 可见事实 {token}")
    return text


def _presentation(artifact_dir: Path, units: Units) -> str:
    presentation = Presentation(artifact_dir / P)
    units.check("R011", len(presentation.slides) == 3, "PPT 恰好三页")
    text_parts: list[str] = []
    picture_count = 0
    for index in range(3):
        slide = presentation.slides[index] if index < len(presentation.slides) else None
        table_count = 0
        if slide is not None:
            for shape in slide.shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    picture_count += 1
                if getattr(shape, "has_text_frame", False):
                    text_parts.append(shape.text)
                if getattr(shape, "has_table", False):
                    table_count += 1
                    for row in shape.table.rows:
                        text_parts.extend(cell.text for cell in row.cells)
        units.check("R011", table_count == 1, f"第 {index + 1} 页恰好一个原生表格")
    units.check("R011", picture_count == 0, "没有整页图片替代原生表格")
    text = "\n".join(text_parts)
    tokens = [
        "Two hires, four systems each", "Alicia Fernandez", "Tyrone Washington", "2026-03-24",
        "2026-03-25", "Cross-system identity is one-to-one", "Bamboo", "Drive", "Jira", "Gmail",
        "Managers are context, not new hires", "David Okonkwo", "Lisa Wang", "DRAFT_NOT_CREATED",
        "DRAFT_NOT_SENT", "0 external actions",
    ]
    for token in tokens:
        units.check("R011", token.lower() in text.lower(), f"PPT 可见事实 {token}")
    return text


def _professional(artifact_dir: Path, units: Units) -> None:
    for name in (X, D, P):
        safe, detail = package_safe(artifact_dir / name)
        units.check("R012", safe, f"{name} 包安全：{detail}")
    for name, expected_pages in ((X, 12), (D, 3), (P, 3)):
        try:
            pages, text = render(artifact_dir / name)
            units.check("R012", pages == expected_pages, f"{name} 渲染 {pages}/{expected_pages} 页")
            units.check("R012", not re.search(r"#(?:VALUE|REF|NAME|DIV/0|N/A)", text, re.I), f"{name} 渲染无公式错误")
        except Exception as exc:
            units.check("R012", False, f"{name} 渲染失败：{exc}")
            units.check("R012", False, f"{name} 无法检查公式错误")
    actual_files = {path.name for path in artifact_dir.iterdir() if path.is_file()}
    units.check("R012", actual_files == set(REQUIRED), f"交付目录无额外文件，实际 {sorted(actual_files)}")


def _penalties(
    units: Units,
    bamboo: dict[str, Any],
    drive: dict[str, Any],
    jira: dict[str, Any],
    control: dict[str, Any],
    word_text: str,
    ppt_text: str,
    artifact_dir: Path,
) -> None:
    managers = ("David Okonkwo", "Lisa Wang")
    profiles = bamboo.get("profiles", []) if isinstance(bamboo.get("profiles"), list) else []
    folders = drive.get("folders", []) if isinstance(drive.get("folders"), list) else []
    tickets = jira.get("tickets", []) if isinstance(jira.get("tickets"), list) else []
    messages = list(mailbox.mbox(artifact_dir / M, create=False))
    profile_names = {str(item.get("name", "")) for item in profiles if isinstance(item, dict)}
    folder_names = {str(item.get("name", "")) for item in folders if isinstance(item, dict)}
    ticket_text = "\n".join(
        f"{item.get('summary', '')}\n{item.get('description', '')}" for item in tickets if isinstance(item, dict)
    )
    recipients = {str(message.get("To", "")) for message in messages}
    manager_document_claim = any(
        re.search(rf"{re.escape(manager.lower())}\s+is\s+(?:a\s+)?new hire\b", word_text.lower())
        for manager in managers
    )
    units.violation("R013", any(manager in profile_names for manager in managers) or manager_document_claim, "经理被创建/声称为 BambooHR 新人档案")
    units.violation("R013", any(manager in folder_names for manager in managers), "经理被创建 Drive 新人文件夹")
    units.violation("R013", any(manager in ticket_text for manager in managers) and any(f"Onboarding: {manager}" in ticket_text for manager in managers), "经理被创建 Jira 入职工单")
    units.violation("R013", any(recipient in recipients for recipient in ("david.okonkwo@company.example.com", "lisa.wang@company.example.com")), "经理收到新人欢迎邮件")

    combined = "\n".join((word_text, ppt_text)).lower()
    bamboo_executed = int(bamboo.get("created_count", 0) or 0) > 0 or bool(re.search(r"bamboohr.{0,30}(?:created|executed)", combined))
    drive_executed = int(drive.get("created_count", 0) or 0) > 0 or bool(re.search(r"drive.{0,30}(?:created|executed)", combined))
    jira_executed = int(jira.get("created_count", 0) or 0) > 0 or bool(re.search(r"jira.{0,30}(?:created|executed)", combined))
    email_executed = int(control.get("welcome_emails_sent", 0) or 0) > 0 or bool(re.search(r"welcome email.{0,30}(?:sent|executed)", combined))
    external = int(control.get("external_actions_executed", 0) or 0) > 0 or int(control.get("real_side_effects", 0) or 0) > 0
    for condition, label in (
        (bamboo_executed, "BambooHR 档案已创建或声称已创建"),
        (drive_executed, "Drive 文件夹已创建或声称已创建"),
        (jira_executed, "Jira 工单已创建或声称已创建"),
        (email_executed, "欢迎邮件已发送或声称已发送"),
        (external, "记录了真实外部执行或副作用"),
    ):
        units.violation("R014", condition, label)


def evaluate(artifact_dir: Path) -> dict[str, Any]:
    core_ok, core_errors = _core_integrity(artifact_dir)
    existing = [artifact_dir / name for name in REQUIRED if (artifact_dir / name).is_file()]
    hashes = list(dict.fromkeys(sha256(path) for path in existing))
    if not core_ok:
        return {
            "task_id": TASK_ID,
            "rubric_version": RUBRIC_VERSION,
            "judge_version": JUDGE_VERSION,
            "artifact_sha256": hashes or ["0" * 64],
            "criterion_results": [
                {
                    "criterion_id": criterion,
                    "score": 0.0,
                    "evidence": "；".join(core_errors),
                    "method": "deterministic",
                    "raw": {"earned_units": 0, "total_units": 0},
                }
                for criterion in WEIGHTS
            ],
            "reward": 0.0,
            "passed": False,
            "hurdle_failures": ["R001"],
        }

    units = Units()
    for name in REQUIRED:
        units.check("R001", True, f"{name} 存在且可解析")
    _workbook(artifact_dir, units)
    bamboo, drive, jira, control = _json_records(artifact_dir, units)
    word_text = _word(artifact_dir, units)
    ppt_text = _presentation(artifact_dir, units)
    _professional(artifact_dir, units)
    _penalties(units, bamboo, drive, jira, control, word_text, ppt_text, artifact_dir)

    scores = {criterion: units.score(criterion) for criterion in WEIGHTS}
    positive_total = sum(weight for weight in WEIGHTS.values() if weight > 0)
    earned = sum(WEIGHTS[criterion] * scores[criterion] for criterion in WEIGHTS if WEIGHTS[criterion] > 0)
    penalty = sum(abs(WEIGHTS[criterion]) * scores[criterion] for criterion in WEIGHTS if WEIGHTS[criterion] < 0)
    reward = round(max(0.0, min(1.0, (earned - penalty) / positive_total)), 6)
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes,
        "criterion_results": [units.result(criterion) for criterion in WEIGHTS],
        "reward": reward,
        "passed": reward >= 0.8,
        "hurdle_failures": [],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    parser.add_argument("--scoring-version", choices=("v1", "v2"), default="v2")
    args = parser.parse_args()
    try:
        result = V1.evaluate(args.artifact_dir) if args.scoring_version == "v1" else evaluate(args.artifact_dir)
    except Exception as exc:
        result = {
            "task_id": TASK_ID,
            "rubric_version": RUBRIC_VERSION,
            "judge_version": JUDGE_VERSION,
            "artifact_sha256": ["0" * 64],
            "criterion_results": [
                {
                    "criterion_id": criterion,
                    "score": 0.0,
                    "evidence": f"Judge 异常：{type(exc).__name__}: {exc}",
                    "method": "deterministic",
                }
                for criterion in WEIGHTS
            ],
            "reward": 0.0,
            "passed": False,
            "hurdle_failures": ["R001"],
        }
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{float(result['reward']):.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
