#!/usr/bin/env python3
"""AutomationBench 5057 的 v2 分层负例与部分完成回放。"""

from __future__ import annotations

import argparse
import json
import mailbox
import shutil
import subprocess
import sys
import tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
PYTHON = sys.executable


def score(case_dir: Path) -> float:
    result = case_dir.parent / f"{case_dir.name}.json"
    reward = case_dir.parent / f"{case_dir.name}.txt"
    subprocess.run(
        [
            PYTHON, str(ROOT / "tests/judge.py"), "--artifact-dir", str(case_dir),
            "--result", str(result), "--reward-file", str(reward), "--scoring-version", "v2",
        ],
        check=True,
        capture_output=True,
        text=True,
        timeout=240,
    )
    return float(reward.read_text(encoding="utf-8"))


def rewrite_mbox(path: Path, mutate) -> None:
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    mutate(messages)
    with path.open("wb") as handle:
        for index, message in enumerate(messages):
            handle.write(f"From x Sun Mar 22 09:00:{index:02d} 2026\n".encode())
            BytesGenerator(handle, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
            handle.write(b"\n")


def make_cases(root: Path) -> dict[str, Path]:
    names = [
        "missing", "corrupt", "unmodified-input-only", "extra-debug", "wrong-count",
        "manager-profile", "wrong-bamboo", "wrong-parent", "missing-folder", "wrong-jira",
        "manager-email", "omit-tyrone-email", "executed", "formula-removed", "word-false",
        "ppt-omit", "manager-everywhere",
    ]
    cases: dict[str, Path] = {}
    for name in names:
        case_dir = root / name
        shutil.copytree(GOLDEN, case_dir)
        cases[name] = case_dir

    (cases["missing"] / "New_Hire_Onboarding_Review.pptx").unlink()
    (cases["corrupt"] / "New_Hire_Onboarding_Audit.xlsx").write_bytes(b"not-an-xlsx")
    for path in list(cases["unmodified-input-only"].iterdir()):
        if path.is_file():
            path.unlink()
    for path in (ROOT / "data/input_files").iterdir():
        if path.is_file():
            shutil.copy2(path, cases["unmodified-input-only"] / path.name)
    (cases["extra-debug"] / "debug.txt").write_text("temporary", encoding="utf-8")

    payload = json.loads((cases["wrong-count"] / "Onboarding_Execution_Control.json").read_text())
    payload["new_hires"] = 3
    (cases["wrong-count"] / "Onboarding_Execution_Control.json").write_text(json.dumps(payload))

    payload = json.loads((cases["manager-profile"] / "BambooHR_Profile_Create_Drafts.json").read_text())
    payload["profiles"][0]["name"] = "David Okonkwo"
    (cases["manager-profile"] / "BambooHR_Profile_Create_Drafts.json").write_text(json.dumps(payload))

    payload = json.loads((cases["wrong-bamboo"] / "BambooHR_Profile_Create_Drafts.json").read_text())
    payload["profiles"][0]["start_date"] = "2026-03-25"
    (cases["wrong-bamboo"] / "BambooHR_Profile_Create_Drafts.json").write_text(json.dumps(payload))

    payload = json.loads((cases["wrong-parent"] / "Drive_Folder_Create_Drafts.json").read_text())
    payload["folders"][0]["parent_id"] = "root"
    (cases["wrong-parent"] / "Drive_Folder_Create_Drafts.json").write_text(json.dumps(payload))

    payload = json.loads((cases["missing-folder"] / "Drive_Folder_Create_Drafts.json").read_text())
    payload["folders"].pop()
    payload["draft_count"] = 1
    (cases["missing-folder"] / "Drive_Folder_Create_Drafts.json").write_text(json.dumps(payload))

    payload = json.loads((cases["wrong-jira"] / "Jira_Onboarding_Ticket_Drafts.json").read_text())
    payload["tickets"][1]["summary"] = "Onboarding: Lisa Wang"
    (cases["wrong-jira"] / "Jira_Onboarding_Ticket_Drafts.json").write_text(json.dumps(payload))

    rewrite_mbox(
        cases["manager-email"] / "New_Hire_Welcome_Drafts.mbox",
        lambda messages: messages[0].replace_header("To", "david.okonkwo@company.example.com"),
    )
    rewrite_mbox(cases["omit-tyrone-email"] / "New_Hire_Welcome_Drafts.mbox", lambda messages: messages.pop())

    payload = json.loads((cases["executed"] / "Onboarding_Execution_Control.json").read_text())
    payload["external_actions_executed"] = 8
    (cases["executed"] / "Onboarding_Execution_Control.json").write_text(json.dumps(payload))

    workbook = load_workbook(cases["formula-removed"] / "New_Hire_Onboarding_Audit.xlsx")
    workbook["Hire_Decisions"]["O2"] = "READY_DRAFT_NOT_EXECUTED"
    workbook.save(cases["formula-removed"] / "New_Hire_Onboarding_Audit.xlsx")

    document = Document(cases["word-false"] / "New_Hire_Onboarding_Brief.docx")
    document.add_paragraph("David Okonkwo is a new hire.")
    document.save(cases["word-false"] / "New_Hire_Onboarding_Brief.docx")

    presentation = Presentation(cases["ppt-omit"] / "New_Hire_Onboarding_Review.pptx")
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                shape.text = shape.text.replace("Tyrone Washington", "Unknown")
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        cell.text = cell.text.replace("Tyrone Washington", "Unknown")
    presentation.save(cases["ppt-omit"] / "New_Hire_Onboarding_Review.pptx")

    payload = json.loads((cases["manager-everywhere"] / "BambooHR_Profile_Create_Drafts.json").read_text())
    payload["profiles"][1]["name"] = "Lisa Wang"
    (cases["manager-everywhere"] / "BambooHR_Profile_Create_Drafts.json").write_text(json.dumps(payload))
    payload = json.loads((cases["manager-everywhere"] / "Drive_Folder_Create_Drafts.json").read_text())
    payload["folders"][1]["name"] = "Lisa Wang"
    (cases["manager-everywhere"] / "Drive_Folder_Create_Drafts.json").write_text(json.dumps(payload))
    payload = json.loads((cases["manager-everywhere"] / "Jira_Onboarding_Ticket_Drafts.json").read_text())
    payload["tickets"][1]["summary"] = "Onboarding: Lisa Wang"
    payload["tickets"][1]["description"] = payload["tickets"][1]["description"].replace("Tyrone Washington", "Lisa Wang")
    (cases["manager-everywhere"] / "Jira_Onboarding_Ticket_Drafts.json").write_text(json.dumps(payload))
    rewrite_mbox(
        cases["manager-everywhere"] / "New_Hire_Welcome_Drafts.mbox",
        lambda messages: messages[1].replace_header("To", "lisa.wang@company.example.com"),
    )
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    with tempfile.TemporaryDirectory(prefix="onboard5057-negative-") as temporary:
        cases = make_cases(Path(temporary))
        results = {name: score(case_dir) for name, case_dir in cases.items()}

    failures: list[str] = []
    for name in ("missing", "corrupt", "unmodified-input-only"):
        if results[name] != 0.0:
            failures.append(f"{name} 应因核心交付不可用为 0，实际 {results[name]}")
    for name in ("extra-debug", "wrong-count", "wrong-bamboo", "wrong-parent", "missing-folder", "omit-tyrone-email", "formula-removed", "ppt-omit"):
        if not 0.85 <= results[name] < 1.0:
            failures.append(f"{name} 应为轻微局部错误分，实际 {results[name]}")
    for name in ("manager-profile", "wrong-jira", "manager-email", "executed", "word-false"):
        if not 0.45 < results[name] < 0.98:
            failures.append(f"{name} 应为中度/严重部分分，实际 {results[name]}")
    if not 0.45 < results["manager-everywhere"] < 0.85:
        failures.append(f"manager-everywhere 应显著低于普通局部错误，实际 {results['manager-everywhere']}")

    partial = [value for value in results.values() if 0.0 < value < 1.0]
    report = {
        "total": len(results),
        "fatal_zero": sum(value == 0.0 for value in results.values()),
        "partial": len(partial),
        "minimum_partial_reward": min(partial),
        "maximum_partial_reward": max(partial),
        "cases": results,
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if failures:
        raise SystemExit("；".join(failures))


if __name__ == "__main__":
    main()
