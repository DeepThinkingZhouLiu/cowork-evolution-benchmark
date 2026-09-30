#!/usr/bin/env python3
import json,mailbox,shutil,subprocess,sys,tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
R=Path(__file__).resolve().parents[1];G=R/"solution/Golden";PY=sys.executable
def score(c):
 r=c.parent/(c.name+".json");w=c.parent/(c.name+".txt");subprocess.run([PY,str(R/"tests/judge.py"),"--artifact-dir",str(c),"--result",str(r),"--reward-file",str(w)],check=True,capture_output=True,text=True,timeout=240);return float(w.read_text())
def rewrite(p,mut):
 b=mailbox.mbox(p,create=False);ms=list(b);b.close();mut(ms)
 with p.open("wb") as h:
  for i,m in enumerate(ms):h.write(f"From x Sun Mar 22 09:00:{i:02d} 2026\n".encode());BytesGenerator(h,policy=SMTP,mangle_from_=True,maxheaderlen=0).flatten(m);h.write(b"\n")
def main():
 with tempfile.TemporaryDirectory(prefix="onboard5057-neg-") as tmp:
  root=Path(tmp);names=["missing","wrong_count","manager_profile","wrong_bamboo","wrong_parent","missing_folder","wrong_jira","manager_email","omit_tyrone","executed","formula_removed","word_false","ppt_omit"];cs=[]
  for n in names:c=root/n;shutil.copytree(G,c);cs.append(c)
  (cs[0]/"New_Hire_Onboarding_Review.pptx").unlink()
  p=json.loads((cs[1]/"Onboarding_Execution_Control.json").read_text());p["new_hires"]=3;(cs[1]/"Onboarding_Execution_Control.json").write_text(json.dumps(p))
  p=json.loads((cs[2]/"BambooHR_Profile_Create_Drafts.json").read_text());p["profiles"][0]["name"]="David Okonkwo";(cs[2]/"BambooHR_Profile_Create_Drafts.json").write_text(json.dumps(p))
  p=json.loads((cs[3]/"BambooHR_Profile_Create_Drafts.json").read_text());p["profiles"][0]["start_date"]="2026-03-25";(cs[3]/"BambooHR_Profile_Create_Drafts.json").write_text(json.dumps(p))
  p=json.loads((cs[4]/"Drive_Folder_Create_Drafts.json").read_text());p["folders"][0]["parent_id"]="root";(cs[4]/"Drive_Folder_Create_Drafts.json").write_text(json.dumps(p))
  p=json.loads((cs[5]/"Drive_Folder_Create_Drafts.json").read_text());p["folders"].pop();p["draft_count"]=1;(cs[5]/"Drive_Folder_Create_Drafts.json").write_text(json.dumps(p))
  p=json.loads((cs[6]/"Jira_Onboarding_Ticket_Drafts.json").read_text());p["tickets"][1]["summary"]="Onboarding: Lisa Wang";(cs[6]/"Jira_Onboarding_Ticket_Drafts.json").write_text(json.dumps(p))
  rewrite(cs[7]/"New_Hire_Welcome_Drafts.mbox",lambda ms:ms[0].replace_header("To","david.okonkwo@company.example.com"))
  rewrite(cs[8]/"New_Hire_Welcome_Drafts.mbox",lambda ms:ms.pop())
  p=json.loads((cs[9]/"Onboarding_Execution_Control.json").read_text());p["external_actions_executed"]=8;(cs[9]/"Onboarding_Execution_Control.json").write_text(json.dumps(p))
  w=load_workbook(cs[10]/"New_Hire_Onboarding_Audit.xlsx");w["Hire_Decisions"]["O2"]="READY_DRAFT_NOT_EXECUTED";w.save(cs[10]/"New_Hire_Onboarding_Audit.xlsx")
  d=Document(cs[11]/"New_Hire_Onboarding_Brief.docx");d.add_paragraph("David Okonkwo is a new hire.");d.save(cs[11]/"New_Hire_Onboarding_Brief.docx")
  p=Presentation(cs[12]/"New_Hire_Onboarding_Review.pptx")
  for s in p.slides:
   for sh in s.shapes:
    if getattr(sh,"has_text_frame",False):sh.text=sh.text.replace("Tyrone Washington","Unknown")
    if getattr(sh,"has_table",False):
     for row in sh.table.rows:
      for cell in row.cells:cell.text=cell.text.replace("Tyrone Washington","Unknown")
  p.save(cs[12]/"New_Hire_Onboarding_Review.pptx")
  x={c.name:score(c) for c in cs};print(json.dumps(x,indent=2));assert all(v==0 for v in x.values()),x
if __name__=="__main__":main()
