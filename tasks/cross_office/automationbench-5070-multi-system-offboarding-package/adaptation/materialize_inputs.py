#!/usr/bin/env python3
"""只静态化 AutomationBench 5070；不是通用造题器。"""
import hashlib,importlib.util,json,shutil,sys,types
from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.worksheet.table import Table,TableStyleInfo
ROOT=Path(__file__).resolve().parents[1];UP=Path(__file__).with_name("upstream-hr-tasks.py");SNAP=Path(__file__).with_name("upstream-task-5070.json");DATA=ROOT/"data/input_files";ENV=ROOT/"environment/input_files";UPSTREAM_SHA="c173a36a03f43096bd122577584aa32f9df379d8760aa46c43be04c85d101779"
def load_task():
 if hashlib.sha256(UP.read_bytes()).hexdigest()!=UPSTREAM_SHA:raise ValueError("上游哈希错误")
 d=types.ModuleType("datasets");d.Dataset=object;sys.modules["datasets"]=d
 for n in ("automationbench","automationbench.domains","automationbench.domains.hr"):sys.modules[n]=types.ModuleType(n)
 noise=types.ModuleType("automationbench.domains.hr._noise");noise.apply_noise=lambda x:x;sys.modules[noise.__name__]=noise
 s=importlib.util.spec_from_file_location("automationbench_5070",UP);m=importlib.util.module_from_spec(s);s.loader.exec_module(m);t=m.get_hr_offboarding_checklist_multi_app_task();assert t["example_id"]==5070 and t["task"]=="hr.offboarding_multi_system";return t
def audit(t):
 i=t["info"];s=i["initial_state"];assert s["meta"]["current_time"]=="2026-03-18T09:00:00Z" and len(i["assertions"])==9
 sheets=s["google_sheets"]["spreadsheets"][0]["worksheets"];assert [x["title"] for x in sheets]==["Week of March 18","Offboarding Processing Convention"] and [len(x["rows"]) for x in sheets]==[3,1]
 r=sheets[0]["rows"];assert [(x["cells"]["Employee"],x["cells"]["Employee ID"],x["cells"]["Last Day"],x["cells"]["PTO Remaining"]) for x in r]==[("Greg Foster","EMP-2001","2026-03-21","5 days"),("Ian Costa","EMP-2004","2026-03-20","0 days"),("Nadia Petrova","EMP-2005","2026-04-05","3 days")]
def style(ws,widths,name):
 for i,w in enumerate(widths,1):ws.column_dimensions[ws.cell(1,i).column_letter].width=w
 ws.freeze_panes="A2";ws.sheet_view.showGridLines=False
 for c in ws[1]:c.fill=PatternFill("solid",fgColor="17365D");c.font=Font(name="Carlito",bold=True,color="FFFFFF");c.alignment=Alignment(horizontal="center",wrap_text=True)
 for row in ws.iter_rows(min_row=2):
  for c in row:c.font=Font(name="Carlito",size=9);c.alignment=Alignment(vertical="top",wrap_text=True)
 t=Table(displayName=name,ref=ws.dimensions);t.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True);ws.add_table(t);ws.page_setup.orientation="landscape";ws.page_setup.fitToWidth=1;ws.sheet_properties.pageSetUpPr.fitToPage=True
def write_xlsx(sheets,p):
 wb=Workbook();wb.remove(wb.active);ws=wb.create_sheet("Week of March 18");h=["Source Row","Employee","Employee ID","Email","Last Day","Type","PTO Remaining","Status"];ws.append(h)
 for r in sheets[0]["rows"]:ws.append([r["row_id"],r["cells"]["Employee"],r["cells"]["Employee ID"],r["cells"]["Email"],date.fromisoformat(r["cells"]["Last Day"]),r["cells"]["Type"],r["cells"]["PTO Remaining"],r["cells"]["Status"]])
 style(ws,[10,26,18,42,18,18,22,18],"OffboardingWeekSource")
 for row in ws.iter_rows(min_row=2):row[4].number_format="yyyy-mm-dd"
 # XLSX 工作表名最长 31 字符；上游标题原文保留在快照，输入表使用明确短名。
 ws=wb.create_sheet("Offboarding Convention");ws.append(["Source Row","Outcome","Tracker Status"])
 for r in sheets[1]["rows"]:ws.append([r["row_id"],r["cells"]["Outcome"],r["cells"]["Tracker Status"]])
 style(ws,[10,64,24],"OffboardingConventionSource");wb.save(p)
def main():
 t=load_task();audit(t);s=t["info"]["initial_state"];DATA.mkdir(parents=True,exist_ok=True);ENV.mkdir(parents=True,exist_ok=True);snap=json.dumps(t,ensure_ascii=False,indent=2)+"\n"
 if SNAP.exists() and SNAP.read_text()!=snap:raise ValueError("单题快照漂移")
 SNAP.write_text(snap);write_xlsx(s["google_sheets"]["spreadsheets"][0]["worksheets"],DATA/"Offboarding_Master.xlsx");(DATA/"Initial_Application_State.json").write_text(json.dumps({"bamboohr":s["bamboohr"],"jira":s["jira"]},indent=2)+"\n");(DATA/"Workflow_Context.json").write_text(json.dumps({"meta":s["meta"],"scope_start":"2026-03-18","scope_end":"2026-03-22","payroll_recipient":"payroll@company.example.com"},indent=2)+"\n")
 for f in DATA.iterdir():
  if f.is_file():shutil.copy2(f,ENV/f.name)
if __name__=="__main__":main()
