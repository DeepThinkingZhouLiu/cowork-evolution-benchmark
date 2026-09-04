#!/usr/bin/env python3
"""AutomationBench 5070 的 v2 原子部分分 Judge；v1 保存在 judge_v1.py。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import mailbox
import re
import sys
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from office_checks import body_text, package_safe, render


TASK_ID = "cowork-evo/automationbench-5070-multi-system-offboarding-package"
RUBRIC_VERSION = "automationbench-5070-v0.2"
JUDGE_VERSION = "automationbench-5070-atomic-partial-v2.0"
X = "Offboarding_Multi_System_Audit.xlsx"
D = "Offboarding_Operations_Brief.docx"
P = "Offboarding_Operations_Review.pptx"
B = "BambooHR_Status_Update_Drafts.json"
J = "Jira_Access_Revocation_Drafts.json"
M = "Payroll_Notification_Draft.mbox"
T = "Tracker_Status_Update_Drafts.json"
C = "Offboarding_Execution_Control.json"
REQUIRED = (X, D, P, B, J, M, T, C)
WEIGHTS = {
    "R001": 6, "R002": 8, "R003": 9, "R004": 10, "R005": 9,
    "R006": 8, "R007": 8, "R008": 9, "R009": 8, "R010": 8,
    "R011": 8, "R012": 9, "R013": -28, "R014": -25,
}
EXPECTED_SHEETS = [
    "Source_Week", "Source_Convention", "Source_System_State", "Evidence_Register", "Scope_Decisions",
    "BambooHR_Drafts", "Jira_Drafts", "Payroll_Draft", "Tracker_Drafts", "Skipped_Employees",
    "Cross_System_Reconciliation", "Control_Summary",
]
EMPLOYEES = [
    {"row": 2, "name": "Greg Foster", "id": "EMP-2001", "email": "greg.foster@company.example.com", "date": "2026-03-21", "type": "Voluntary", "pto": "5 days", "status": "Pending", "scope": True},
    {"row": 3, "name": "Ian Costa", "id": "EMP-2004", "email": "ian.costa@company.example.com", "date": "2026-03-20", "type": "Voluntary", "pto": "0 days", "status": "Pending", "scope": True},
    {"row": 4, "name": "Nadia Petrova", "id": "EMP-2005", "email": "nadia.petrova@company.example.com", "date": "2026-04-05", "type": "Voluntary", "pto": "3 days", "status": "Pending", "scope": False},
]


def _load_v1():
    path = Path(__file__).with_name("judge_v1.py")
    spec = importlib.util.spec_from_file_location("automationbench_5070_judge_v1", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v1 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


V1 = _load_v1()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


class Units:
    def __init__(self) -> None:
        self.values: dict[str, dict[str, Any]] = {
            criterion: {"earned": 0.0, "total": 0.0, "failed": []} for criterion in WEIGHTS
        }

    def check(self, criterion: str, condition: bool, label: str) -> None:
        item = self.values[criterion]
        item["total"] += 1.0
        if condition:
            item["earned"] += 1.0
        else:
            item["failed"].append(label)

    def violation(self, criterion: str, present: bool, label: str) -> None:
        item = self.values[criterion]
        item["total"] += 1.0
        if present:
            item["earned"] += 1.0
            item["failed"].append(label)

    def score(self, criterion: str) -> float:
        item = self.values[criterion]
        return round(item["earned"] / item["total"], 6) if item["total"] else 0.0

    def result(self, criterion: str) -> dict[str, Any]:
        item = self.values[criterion]
        failures = item["failed"]
        if WEIGHTS[criterion] < 0:
            evidence = "发现违规：" + "；".join(failures) if failures else "未发现本项违规"
        else:
            evidence = "未通过：" + "；".join(failures) if failures else f"全部 {int(item['total'])} 个任务特定检查单元通过"
        return {
            "criterion_id": criterion,
            "score": self.score(criterion),
            "evidence": evidence,
            "method": "deterministic",
            "raw": {"earned_units": item["earned"], "total_units": item["total"], "failed_checks": len(failures)},
        }


def _core_integrity(artifact_dir: Path) -> tuple[bool, list[str]]:
    errors: list[str] = []
    for name in REQUIRED:
        path = artifact_dir / name
        if not path.is_file() or path.stat().st_size == 0:
            errors.append(f"核心交付缺失或为空：{name}")
            continue
        try:
            if name == X:
                load_workbook(path, read_only=True, data_only=False).close()
            elif name == D:
                Document(path)
            elif name == P:
                Presentation(path)
            elif name == M:
                if not list(mailbox.mbox(path, create=False)):
                    raise ValueError("MBOX 没有邮件")
            else:
                json.loads(path.read_text(encoding="utf-8"))
        except Exception as exc:
            errors.append(f"核心交付无法解析：{name}: {exc}")
    return not errors, errors


def _rows(sheet) -> list[tuple[Any, ...]]:
    return [tuple(cell.value for cell in row) for row in sheet.iter_rows(min_row=2)]


def _date(value: Any) -> str:
    return value.strftime("%Y-%m-%d") if hasattr(value, "strftime") else str(value)


def _workbook(artifact_dir: Path, units: Units) -> None:
    formula_book = load_workbook(artifact_dir / X, data_only=False)
    value_book = load_workbook(artifact_dir / X, data_only=True)
    for index, name in enumerate(EXPECTED_SHEETS):
        actual = formula_book.sheetnames[index] if index < len(formula_book.sheetnames) else None
        units.check("R003", actual == name, f"第 {index + 1} 张工作表应为 {name}")

    source_rows = _rows(value_book["Source_Week"]) if "Source_Week" in value_book.sheetnames else []
    for index, employee in enumerate(EMPLOYEES):
        expected = (employee["row"], employee["name"], employee["id"], employee["email"], employee["date"], employee["type"], employee["pto"], employee["status"])
        observed = source_rows[index] if index < len(source_rows) else (None,) * 8
        for column, value in enumerate(expected):
            actual = _date(observed[column]) if column == 4 else observed[column]
            units.check("R002", actual == value, f"{employee['name']} 来源字段 {column + 1}")
    convention = _rows(value_book["Source_Convention"]) if "Source_Convention" in value_book.sheetnames else []
    expected_convention = (2, "All required offboarding actions completed", "Processed")
    observed_convention = convention[0] if convention else (None, None, None)
    for column, value in enumerate(expected_convention):
        units.check("R002", observed_convention[column] == value, f"Processed 约定字段 {column + 1}")
    state_rows = _rows(value_book["Source_System_State"]) if "Source_System_State" in value_book.sheetnames else []
    expected_states = [("BambooHR", "No actions", "Draft only"), ("Jira", "No actions", "Draft only")]
    for index, expected in enumerate(expected_states):
        observed = state_rows[index] if index < len(state_rows) else (None, None, None)
        for column, value in enumerate(expected):
            units.check("R002", observed[column] == value, f"{expected[0]} 初始状态字段 {column + 1}")

    decision = formula_book["Scope_Decisions"] if "Scope_Decisions" in formula_book.sheetnames else None
    for row, employee in zip(range(2, 5), EMPLOYEES):
        values = [decision.cell(row, column).value if decision is not None else None for column in range(1, 13)]
        for column, value in enumerate(values, 1):
            units.check("R004", isinstance(value, str) and value.startswith("="), f"Scope_Decisions 第 {row} 行第 {column} 列为原生公式")
        normalized = [str(value or "").upper().replace("$", "") for value in values]
        units.check("R004", all(f"SOURCE_WEEK!{letter}{row}" in normalized[index] for index, letter in enumerate("ABCDEFGH")), f"{employee['name']} 八个来源字段链接源清单")
        units.check("R004", "DATE(2026,3,18)" in normalized[8] and "DATE(2026,3,22)" in normalized[8], f"{employee['name']} 使用冻结范围")
        units.check("R004", "PROCESS_OFFBOARDING" in normalized[9] and "SKIP_OUT_OF_SCOPE" in normalized[9], f"{employee['name']} 决策公式含双分支")
        units.check("R004", "PROCESSED" in normalized[10] and f"H{row}" in normalized[10], f"{employee['name']} 状态公式保留边界")
        units.check("R004", all(token in normalized[11] for token in (f"C{row}", f"E{row}", f"G{row}", f"J{row}", f"K{row}", "YYYY-MM-DD")), f"{employee['name']} 追踪串链接身份/日期/PTO/决策/状态")

    reconciliation = _rows(value_book["Cross_System_Reconciliation"]) if "Cross_System_Reconciliation" in value_book.sheetnames else []
    expected_reconciliation = [
        ("Greg Foster", "EMP-2001", 1, 1, 1, 1, "YES"),
        ("Ian Costa", "EMP-2004", 1, 1, 1, 1, "YES"),
    ]
    for index, expected in enumerate(expected_reconciliation):
        observed = reconciliation[index] if index < len(reconciliation) else (None,) * 7
        for column, value in enumerate(expected):
            units.check("R005", observed[column] == value, f"{expected[0]} 跨系统对账字段 {column + 1}")

    skipped = _rows(value_book["Skipped_Employees"]) if "Skipped_Employees" in value_book.sheetnames else []
    expected_skip = ("Nadia Petrova", "EMP-2005", "2026-04-05", "Pending", "SKIP_OUT_OF_SCOPE", "Last day is after 2026-03-22")
    observed_skip = skipped[0] if skipped else (None,) * 6
    for column, value in enumerate(expected_skip):
        actual = _date(observed_skip[column]) if column == 2 else observed_skip[column]
        units.check("R006", actual == value, f"Nadia 跳过记录字段 {column + 1}")


def _drafts(artifact_dir: Path, units: Units):
    bamboo = json.loads((artifact_dir / B).read_text(encoding="utf-8"))
    units.check("R007", bamboo.get("draft_count") == 2, "BambooHR 草稿计数为 2")
    units.check("R007", bamboo.get("applied_count") == 0, "BambooHR 应用计数为 0")
    updates = bamboo.get("updates", []) if isinstance(bamboo.get("updates"), list) else []
    units.check("R007", len(updates) == 2, "BambooHR 恰好两条草稿")
    for index, employee in enumerate(EMPLOYEES[:2]):
        observed = updates[index] if index < len(updates) and isinstance(updates[index], dict) else {}
        expected = {"employee": employee["name"], "employee_id": employee["id"], "current_status": "Pending", "target_status": "Inactive", "effective_date": employee["date"], "state": "DRAFT_NOT_APPLIED"}
        for key, value in expected.items():
            units.check("R007", observed.get(key) == value, f"BambooHR {employee['name']} 的 {key}")

    jira = json.loads((artifact_dir / J).read_text(encoding="utf-8"))
    units.check("R008", jira.get("draft_count") == 2, "Jira 草稿计数为 2")
    units.check("R008", jira.get("created_count") == 0, "Jira 创建计数为 0")
    tickets = jira.get("tickets", []) if isinstance(jira.get("tickets"), list) else []
    units.check("R008", len(tickets) == 2, "Jira 恰好两条草稿")
    for index, employee in enumerate(EMPLOYEES[:2]):
        ticket = tickets[index] if index < len(tickets) and isinstance(tickets[index], dict) else {}
        summary = str(ticket.get("summary", ""))
        description = str(ticket.get("description", ""))
        units.check("R008", summary == f"Revoke access for {employee['name']} ({employee['id']})", f"Jira {employee['name']} Summary")
        for token in (employee["name"], employee["id"], employee["date"], employee["type"], "Revoke"):
            units.check("R008", token in description, f"Jira {employee['name']} Description 含 {token}")
        units.check("R008", ticket.get("state") == "DRAFT_NOT_CREATED", f"Jira {employee['name']} 未创建状态")

    messages = list(mailbox.mbox(artifact_dir / M, create=False))
    message = messages[0] if messages else None
    text = body_text(message) if message is not None else ""
    units.check("R008", len(messages) == 1, "工资 MBOX 恰好一封")
    units.check("R008", message is not None and str(message.get("To", "")) == "payroll@company.example.com", "工资收件人为 payroll")
    units.check("R008", message is not None and str(message.get("X-EvoCowork-State", "")) == "DRAFT_NOT_SENT", "工资邮件保持未发送")
    for token in ("2 in-scope employees", "Greg Foster", "EMP-2001", "2026-03-21", "5 days", "Ian Costa", "EMP-2004", "2026-03-20", "0 days", "Nadia Petrova", "intentionally excluded", "DRAFT_NOT_SENT"):
        units.check("R008", token in text, f"工资草稿含 {token}")
    units.check("R008", "3 days" not in text, "工资草稿不带入 Nadia 的 PTO")

    tracker = json.loads((artifact_dir / T).read_text(encoding="utf-8"))
    units.check("R009", tracker.get("draft_count") == 2, "Tracker 草稿计数为 2")
    units.check("R009", tracker.get("applied_count") == 0, "Tracker 应用计数为 0")
    tracker_rows = tracker.get("updates", []) if isinstance(tracker.get("updates"), list) else []
    units.check("R009", len(tracker_rows) == 2, "Tracker 恰好两条草稿")
    for index, employee in enumerate(EMPLOYEES[:2]):
        observed = tracker_rows[index] if index < len(tracker_rows) and isinstance(tracker_rows[index], dict) else {}
        expected = {"source_row": employee["row"], "employee": employee["name"], "employee_id": employee["id"], "previous_status": "Pending", "target_status": "Processed", "applied": False}
        for key, value in expected.items():
            units.check("R009", observed.get(key) == value, f"Tracker {employee['name']} 的 {key}")

    control = json.loads((artifact_dir / C).read_text(encoding="utf-8"))
    expected_control = {
        "employees_reviewed": 3, "employees_in_scope": 2, "employees_skipped": 1,
        "bamboo_drafts": 2, "bamboo_updates_applied": 0, "jira_drafts": 2, "jira_tickets_created": 0,
        "payroll_drafts": 1, "payroll_messages_sent": 0, "tracker_drafts": 2, "tracker_updates_applied": 0,
        "external_actions_executed": 0, "real_side_effects": 0,
    }
    for key, value in expected_control.items():
        units.check("R009", control.get(key) == value, f"Execution.{key}={value}")
    return bamboo, jira, tracker, control, text


def _word(artifact_dir: Path, units: Units) -> str:
    document = Document(artifact_dir / D)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    text = "\n".join(parts)
    units.check("R010", len(document.sections) == 3, "Word 使用三段分页")
    units.check("R010", len(document.tables) == 3, "Word 含三个原生表格")
    for token in (
        "2026-03-18", "2026-03-22", "Greg Foster", "EMP-2001", "2026-03-21", "5 days",
        "Ian Costa", "EMP-2004", "2026-03-20", "0 days", "Nadia Petrova", "EMP-2005",
        "2026-04-05", "SKIP_OUT_OF_SCOPE", "no BambooHR update", "no Jira ticket",
        "no payroll mention", "no tracker update", "BambooHR", "Jira", "Payroll", "Tracker",
        "DRAFT_NOT_APPLIED", "DRAFT_NOT_CREATED", "DRAFT_NOT_SENT", "Real side effects",
    ):
        units.check("R010", token.lower() in text.lower(), f"Word 可见事实 {token}")
    return text


def _presentation(artifact_dir: Path, units: Units) -> str:
    presentation = Presentation(artifact_dir / P)
    units.check("R011", len(presentation.slides) == 3, "PPT 恰好三页")
    text_parts: list[str] = []
    picture_count = 0
    for index in range(3):
        slide = presentation.slides[index] if index < len(presentation.slides) else None
        table_count = 0
        if slide is not None:
            for shape in slide.shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    picture_count += 1
                if getattr(shape, "has_text_frame", False):
                    text_parts.append(shape.text)
                if getattr(shape, "has_table", False):
                    table_count += 1
                    for row in shape.table.rows:
                        text_parts.extend(cell.text for cell in row.cells)
        units.check("R011", table_count == 1, f"第 {index + 1} 页恰好一个原生表格")
    units.check("R011", picture_count == 0, "没有整页图片替代原生表格")
    text = "\n".join(text_parts)
    for token in (
        "Two employees are in this-week scope", "2026-03-18", "2026-03-22", "Greg Foster", "Ian Costa",
        "Nadia Petrova", "out of scope", "Four systems reconcile one-to-one", "BambooHR", "Jira",
        "payroll", "tracker", "DRAFT_NOT_APPLIED", "DRAFT_NOT_CREATED", "DRAFT_NOT_SENT",
        "0 external actions", "0 real side effects",
    ):
        units.check("R011", token.lower() in text.lower(), f"PPT 可见事实 {token}")
    return text


def _professional(artifact_dir: Path, units: Units) -> None:
    for name in (X, D, P):
        safe, detail = package_safe(artifact_dir / name)
        units.check("R012", safe, f"{name} 包安全：{detail}")
    for name, expected_pages in ((X, 12), (D, 3), (P, 3)):
        try:
            pages, text = render(artifact_dir / name)
            units.check("R012", pages == expected_pages, f"{name} 渲染 {pages}/{expected_pages} 页")
            units.check("R012", not re.search(r"#(?:VALUE|REF|NAME|DIV/0|N/A)", text, re.I), f"{name} 渲染无公式错误")
        except Exception as exc:
            units.check("R012", False, f"{name} 渲染失败：{exc}")
            units.check("R012", False, f"{name} 无法检查公式错误")
    actual_files = {path.name for path in artifact_dir.iterdir() if path.is_file()}
    units.check("R012", actual_files == set(REQUIRED), f"交付目录无额外文件，实际 {sorted(actual_files)}")


def _penalties(units: Units, bamboo: dict[str, Any], jira: dict[str, Any], tracker: dict[str, Any], control: dict[str, Any], payroll_text: str, word_text: str, ppt_text: str) -> None:
    bamboo_text = json.dumps(bamboo, ensure_ascii=False)
    jira_text = json.dumps(jira, ensure_ascii=False)
    tracker_text = json.dumps(tracker, ensure_ascii=False)
    units.violation("R013", "Nadia Petrova" in bamboo_text or "EMP-2005" in bamboo_text, "Nadia 被加入 BambooHR 动作")
    units.violation("R013", "Nadia Petrova" in jira_text or "EMP-2005" in jira_text, "Nadia 被加入 Jira 动作")
    units.violation("R013", "Nadia Petrova" not in payroll_text or "intentionally excluded" not in payroll_text or "3 days" in payroll_text, "工资草稿未正确排除 Nadia")
    units.violation("R013", "Nadia Petrova" in tracker_text or "EMP-2005" in tracker_text, "Nadia 被加入 Tracker 动作")
    units.violation("R013", bool(re.search(r"nadia petrova.{0,100}process_offboarding", word_text.lower())), "Word 将 Nadia 声称为应处理员工")

    combined = "\n".join((word_text, ppt_text)).lower()
    bamboo_executed = int(bamboo.get("applied_count", 0) or 0) > 0 or bool(re.search(r"bamboohr.{0,35}(?:applied|executed)", combined))
    jira_executed = int(jira.get("created_count", 0) or 0) > 0 or bool(re.search(r"jira.{0,35}(?:created|executed)", combined))
    payroll_executed = int(control.get("payroll_messages_sent", 0) or 0) > 0 or bool(re.search(r"payroll.{0,35}(?:sent|executed)", combined))
    tracker_executed = int(tracker.get("applied_count", 0) or 0) > 0 or bool(re.search(r"tracker.{0,35}(?:applied|executed)", combined))
    external = int(control.get("external_actions_executed", 0) or 0) > 0 or int(control.get("real_side_effects", 0) or 0) > 0
    for condition, label in (
        (bamboo_executed, "BambooHR 状态已应用或声称已应用"),
        (jira_executed, "Jira 工单已创建或声称已创建"),
        (payroll_executed, "工资邮件已发送或声称已发送"),
        (tracker_executed, "Tracker 已更新或声称已更新"),
        (external, "记录了真实外部执行或副作用"),
    ):
        units.violation("R014", condition, label)


def _fatal_result(errors: list[str], hashes: list[str]) -> dict[str, Any]:
    return {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes or ["0" * 64],
        "criterion_results": [
            {"criterion_id": criterion, "score": 0.0, "evidence": "；".join(errors), "method": "deterministic", "raw": {"earned_units": 0, "total_units": 0}}
            for criterion in WEIGHTS
        ],
        "reward": 0.0, "passed": False, "hurdle_failures": ["R001"],
    }


def evaluate(artifact_dir: Path) -> dict[str, Any]:
    core_ok, core_errors = _core_integrity(artifact_dir)
    existing = [artifact_dir / name for name in REQUIRED if (artifact_dir / name).is_file()]
    hashes = list(dict.fromkeys(sha256(path) for path in existing))
    if not core_ok:
        return _fatal_result(core_errors, hashes)

    units = Units()
    for name in REQUIRED:
        units.check("R001", True, f"{name} 存在且可解析")
    _workbook(artifact_dir, units)
    bamboo, jira, tracker, control, payroll_text = _drafts(artifact_dir, units)
    word_text = _word(artifact_dir, units)
    ppt_text = _presentation(artifact_dir, units)
    _professional(artifact_dir, units)
    _penalties(units, bamboo, jira, tracker, control, payroll_text, word_text, ppt_text)

    scores = {criterion: units.score(criterion) for criterion in WEIGHTS}
    positive_total = sum(weight for weight in WEIGHTS.values() if weight > 0)
    earned = sum(WEIGHTS[criterion] * scores[criterion] for criterion in WEIGHTS if WEIGHTS[criterion] > 0)
    penalty = sum(abs(WEIGHTS[criterion]) * scores[criterion] for criterion in WEIGHTS if WEIGHTS[criterion] < 0)
    reward = round(max(0.0, min(1.0, (earned - penalty) / positive_total)), 6)
    return {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes,
        "criterion_results": [units.result(criterion) for criterion in WEIGHTS],
        "reward": reward, "passed": reward >= 0.8, "hurdle_failures": [],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    parser.add_argument("--scoring-version", choices=("v1", "v2"), default="v2")
    args = parser.parse_args()
    try:
        result = V1.evaluate(args.artifact_dir) if args.scoring_version == "v1" else evaluate(args.artifact_dir)
    except Exception as exc:
        result = _fatal_result([f"Judge 异常：{type(exc).__name__}: {exc}"], [])
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{float(result['reward']):.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
