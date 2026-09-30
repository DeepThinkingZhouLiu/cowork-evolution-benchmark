#!/usr/bin/env python3
"""5070 v2 专属分层负例：范围、身份、PTO、公式、对账与离线执行边界。"""

from __future__ import annotations

import argparse
import json
import mailbox
import shutil
import subprocess
import sys
import tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from typing import Callable

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution/Golden"
PYTHON = sys.executable


def score(case_dir: Path) -> float:
    result = case_dir.parent / f"{case_dir.name}.result.json"
    reward = case_dir.parent / f"{case_dir.name}.reward.txt"
    subprocess.run(
        [PYTHON, str(ROOT / "tests/judge.py"), "--artifact-dir", str(case_dir), "--result", str(result), "--reward-file", str(reward), "--scoring-version", "v2"],
        check=True, capture_output=True, text=True, timeout=300,
    )
    return float(reward.read_text())


def rewrite_mbox(path: Path, transform: Callable[[list], None]) -> None:
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    transform(messages)
    with path.open("wb") as handle:
        for index, message in enumerate(messages):
            handle.write(f"From x Wed Mar 18 09:00:{index:02d} 2026\n".encode())
            BytesGenerator(handle, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
            handle.write(b"\n")


def replace_docx(path: Path, old: str, new: str) -> None:
    document = Document(path)
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.text = run.text.replace(old, new)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.text = run.text.replace(old, new)
    document.save(path)


def replace_pptx(path: Path, old: str, new: str) -> None:
    presentation = Presentation(path)
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.text = run.text.replace(old, new)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        cell.text = cell.text.replace(old, new)
    presentation.save(path)


def make_cases(root: Path) -> dict[str, Path]:
    names = [
        "missing", "corrupt", "unmodified-input-only", "extra-debug", "wrong-count", "nadia-bamboo",
        "wrong-effective-date", "nadia-jira", "wrong-pto", "omit-ian-payroll", "wrong-tracker",
        "executed", "formula-removed", "wrong-reconciliation", "word-scope-false", "ppt-omit-nadia",
        "nadia-everywhere",
    ]
    cases: dict[str, Path] = {}
    for name in names:
        case_dir = root / name
        shutil.copytree(GOLDEN, case_dir)
        cases[name] = case_dir

    (cases["missing"] / "Offboarding_Operations_Review.pptx").unlink()
    (cases["corrupt"] / "Offboarding_Multi_System_Audit.xlsx").write_bytes(b"not-an-xlsx")
    for path in list(cases["unmodified-input-only"].iterdir()):
        if path.is_file():
            path.unlink()
    for path in (ROOT / "data/input_files").iterdir():
        if path.is_file():
            shutil.copy2(path, cases["unmodified-input-only"] / path.name)
    (cases["extra-debug"] / "debug.txt").write_text("temporary", encoding="utf-8")

    payload = json.loads((cases["wrong-count"] / "Offboarding_Execution_Control.json").read_text())
    payload["employees_in_scope"] = 3
    (cases["wrong-count"] / "Offboarding_Execution_Control.json").write_text(json.dumps(payload))

    payload = json.loads((cases["nadia-bamboo"] / "BambooHR_Status_Update_Drafts.json").read_text())
    payload["updates"][1].update({"employee": "Nadia Petrova", "employee_id": "EMP-2005", "effective_date": "2026-04-05"})
    (cases["nadia-bamboo"] / "BambooHR_Status_Update_Drafts.json").write_text(json.dumps(payload))

    payload = json.loads((cases["wrong-effective-date"] / "BambooHR_Status_Update_Drafts.json").read_text())
    payload["updates"][0]["effective_date"] = "2026-03-22"
    (cases["wrong-effective-date"] / "BambooHR_Status_Update_Drafts.json").write_text(json.dumps(payload))

    payload = json.loads((cases["nadia-jira"] / "Jira_Access_Revocation_Drafts.json").read_text())
    payload["tickets"][1]["summary"] = "Revoke access for Nadia Petrova (EMP-2005)"
    payload["tickets"][1]["description"] = "Revoke company-system access for Nadia Petrova (EMP-2005) after the last day 2026-04-05."
    (cases["nadia-jira"] / "Jira_Access_Revocation_Drafts.json").write_text(json.dumps(payload))

    rewrite_mbox(
        cases["wrong-pto"] / "Payroll_Notification_Draft.mbox",
        lambda messages: messages[0].set_payload(messages[0].get_payload().replace("5 da=\nys", "6 da=\nys").replace("5 days", "6 days")),
    )
    rewrite_mbox(
        cases["omit-ian-payroll"] / "Payroll_Notification_Draft.mbox",
        lambda messages: messages[0].set_payload(messages[0].get_payload().replace("Ian Costa", "Unknown Employee")),
    )

    payload = json.loads((cases["wrong-tracker"] / "Tracker_Status_Update_Drafts.json").read_text())
    payload["updates"][0]["target_status"] = "Pending"
    (cases["wrong-tracker"] / "Tracker_Status_Update_Drafts.json").write_text(json.dumps(payload))

    payload = json.loads((cases["executed"] / "Offboarding_Execution_Control.json").read_text())
    payload["external_actions_executed"] = 7
    payload["real_side_effects"] = 7
    (cases["executed"] / "Offboarding_Execution_Control.json").write_text(json.dumps(payload))

    workbook = load_workbook(cases["formula-removed"] / "Offboarding_Multi_System_Audit.xlsx")
    workbook["Scope_Decisions"]["J2"] = "PROCESS_OFFBOARDING"
    workbook.save(cases["formula-removed"] / "Offboarding_Multi_System_Audit.xlsx")

    workbook = load_workbook(cases["wrong-reconciliation"] / "Offboarding_Multi_System_Audit.xlsx")
    workbook["Cross_System_Reconciliation"]["G2"] = "NO"
    workbook.save(cases["wrong-reconciliation"] / "Offboarding_Multi_System_Audit.xlsx")

    replace_docx(cases["word-scope-false"] / "Offboarding_Operations_Brief.docx", "SKIP_OUT_OF_SCOPE", "PROCESS_OFFBOARDING")
    replace_pptx(cases["ppt-omit-nadia"] / "Offboarding_Operations_Review.pptx", "Nadia Petrova", "Unknown Employee")

    payload = json.loads((cases["nadia-everywhere"] / "BambooHR_Status_Update_Drafts.json").read_text())
    payload["updates"][1].update({"employee": "Nadia Petrova", "employee_id": "EMP-2005", "effective_date": "2026-04-05"})
    (cases["nadia-everywhere"] / "BambooHR_Status_Update_Drafts.json").write_text(json.dumps(payload))
    payload = json.loads((cases["nadia-everywhere"] / "Jira_Access_Revocation_Drafts.json").read_text())
    payload["tickets"][1] = {"summary": "Revoke access for Nadia Petrova (EMP-2005)", "description": "Revoke Nadia Petrova EMP-2005 after 2026-04-05.", "state": "DRAFT_NOT_CREATED"}
    (cases["nadia-everywhere"] / "Jira_Access_Revocation_Drafts.json").write_text(json.dumps(payload))
    payload = json.loads((cases["nadia-everywhere"] / "Tracker_Status_Update_Drafts.json").read_text())
    payload["updates"][1].update({"source_row": 4, "employee": "Nadia Petrova", "employee_id": "EMP-2005"})
    (cases["nadia-everywhere"] / "Tracker_Status_Update_Drafts.json").write_text(json.dumps(payload))
    rewrite_mbox(
        cases["nadia-everywhere"] / "Payroll_Notification_Draft.mbox",
        lambda messages: messages[0].set_payload(messages[0].get_payload().replace("intentionally excluded", "included with 3 days PTO")),
    )
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    with tempfile.TemporaryDirectory(prefix="offboarding5070-negative-") as temporary:
        cases = make_cases(Path(temporary))
        results = {name: score(case_dir) for name, case_dir in cases.items()}

    failures: list[str] = []
    for name in ("missing", "corrupt", "unmodified-input-only"):
        if results[name] != 0.0:
            failures.append(f"{name} 应因核心交付不可用为 0，实际 {results[name]}")
    for name in ("extra-debug", "wrong-count", "wrong-effective-date", "wrong-pto", "omit-ian-payroll", "wrong-tracker", "formula-removed", "wrong-reconciliation", "ppt-omit-nadia"):
        if not 0.82 <= results[name] < 1.0:
            failures.append(f"{name} 应为局部错误部分分，实际 {results[name]}")
    for name in ("nadia-bamboo", "nadia-jira", "executed", "word-scope-false"):
        if not 0.35 < results[name] < 0.96:
            failures.append(f"{name} 应为中度/严重部分分，实际 {results[name]}")
    if not 0.05 < results["nadia-everywhere"] < 0.75:
        failures.append(f"nadia-everywhere 应显著低于普通局部错误，实际 {results['nadia-everywhere']}")

    partial = [value for value in results.values() if 0.0 < value < 1.0]
    report = {
        "total": len(results), "fatal_zero": sum(value == 0.0 for value in results.values()),
        "partial": len(partial), "minimum_partial_reward": min(partial), "maximum_partial_reward": max(partial),
        "cases": results,
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if failures:
        raise SystemExit("；".join(failures))


if __name__ == "__main__":
    main()
