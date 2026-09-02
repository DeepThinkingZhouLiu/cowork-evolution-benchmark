#!/usr/bin/env python3
"""只把 AutomationBench 5090 的公开初始状态静态化为本题输入。

本脚本只接受 ``hr.jira_security_access_review``，并逐项断言五名员工、
三封证据邮件、公开频道、运行时间和十二条上游断言。它不是通用制题器。
"""

from __future__ import annotations

import importlib.util
import json
import shutil
import sys
import types
import zipfile
from datetime import datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
UPSTREAM = Path(__file__).with_name("upstream-hr-tasks.py")
SNAPSHOT = Path(__file__).with_name("upstream-task-5090.json")
DATA_DIR = ROOT / "data" / "input_files"
ENV_DIR = ROOT / "environment" / "input_files"
FIXED_ZIP_TIME = (2026, 4, 7, 9, 0, 0)


def load_task() -> dict:
    datasets = types.ModuleType("datasets")
    datasets.Dataset = object
    sys.modules["datasets"] = datasets
    for name in ("automationbench", "automationbench.domains", "automationbench.domains.hr"):
        sys.modules[name] = types.ModuleType(name)
    noise = types.ModuleType("automationbench.domains.hr._noise")
    noise.apply_noise = lambda value: value
    sys.modules[noise.__name__] = noise

    spec = importlib.util.spec_from_file_location("automationbench_hr_tasks_frozen", UPSTREAM)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入冻结的 AutomationBench HR 源文件")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    task = module.get_hr_jira_security_access_review_task()
    if task.get("example_id") != 5090 or task.get("task") != "hr.jira_security_access_review":
        raise ValueError("冻结函数不再是 AutomationBench 5090")
    return task


def assert_shape(task: dict) -> None:
    state = task["info"]["initial_state"]
    assert state["meta"] == {"schema_version": "0.1.0", "current_time": "2026-04-07T09:00:00Z"}
    assert [item["id"] for item in state["gmail"]["messages"]] == [
        "msg_access_review_policy", "msg_dept_correction", "msg_termination_notice"
    ]
    source = state["google_sheets"]["spreadsheets"]
    assert len(source) == 1 and source[0]["id"] == "ss_access_review"
    rows = source[0]["worksheets"][0]["rows"]
    assert [item["row_id"] for item in rows] == [2, 3, 4, 5, 6]
    assert [item["cells"]["Employee"] for item in rows] == [
        "Alice Park", "Bob Chen", "Sophie Laurent", "Wei Zhang", "Carol Diaz"
    ]
    assert state["slack"]["channels"] == [
        {"id": "C_ANN_5090", "name": "company-announcements", "is_private": False}
    ]
    assert state["slack"]["messages"] == []
    assert state["jira"] == {"actions": {}}
    assert len(task["info"]["assertions"]) == 12


def normalize_xlsx(path: Path) -> None:
    temporary = path.with_suffix(".normalized.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        temporary, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as destination:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            destination.writestr(info, source.read(name))
    temporary.replace(path)


def write_tracker(state: dict, target: Path) -> None:
    source = state["google_sheets"]["spreadsheets"][0]
    worksheet = source["worksheets"][0]
    headers = [
        "Source Row", "Employee", "Employee ID", "Department", "Manager", "Manager Email",
        "Access Level", "Role Level", "Flag", "Status", "Notes",
    ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = worksheet["title"]
    workbook.creator = "AutomationBench 5090 public fixture"
    workbook.title = source["title"]
    workbook.subject = "Static Office representation of the Q2 access review source"
    workbook.description = "No employee, access level, role, flag or status was added beyond AutomationBench 5090."
    sheet.append(headers)
    for item in worksheet["rows"]:
        cells = item["cells"]
        sheet.append([item["row_id"]] + [cells.get(header, "") for header in headers[1:]])
    table = Table(displayName="AccessAuditSourceInput", ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    widths = {
        "A": 12, "B": 22, "C": 16, "D": 18, "E": 22, "F": 34,
        "G": 16, "H": 14, "I": 16, "J": 14, "K": 28,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = sheet.dimensions
    workbook.save(target)
    normalize_xlsx(target)


def message_bytes(record: dict) -> bytes:
    message = EmailMessage(policy=SMTP)
    message["From"] = record["from_"]
    message["To"] = ", ".join(record["to"])
    message["Subject"] = record["subject"]
    parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
    message["Date"] = parsed.strftime("%a, %d %b %Y %H:%M:%S +0000")
    message["Message-ID"] = f"<{record['id']}@automationbench.invalid>"
    message["X-AutomationBench-Message-ID"] = record["id"]
    message["X-AutomationBench-Thread-ID"] = record["thread_id"]
    message["X-AutomationBench-Labels"] = ",".join(record.get("label_ids", []))
    message.set_content(record["body_plain"], charset="utf-8")
    buffer = BytesIO()
    BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
    return buffer.getvalue()


def write_mbox(messages: list[dict], target: Path) -> None:
    chunks: list[bytes] = []
    for record in messages:
        parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
        chunks.append(f"From {record['from_']} {parsed.strftime('%a %b %d %H:%M:%S %Y')}\n".encode("ascii"))
        chunks.append(message_bytes(record).replace(b"\r\n", b"\n"))
        chunks.append(b"\n")
    target.write_bytes(b"".join(chunks))


def main() -> None:
    task = load_task()
    assert_shape(task)
    state = task["info"]["initial_state"]
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    ENV_DIR.mkdir(parents=True, exist_ok=True)

    SNAPSHOT.write_text(json.dumps(task, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_tracker(state, DATA_DIR / "Q2_Access_Review.xlsx")
    write_mbox(state["gmail"]["messages"], DATA_DIR / "Access_Review_Mailbox.mbox")
    (DATA_DIR / "Company_Announcements_Channel.json").write_text(
        json.dumps(state["slack"], ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (DATA_DIR / "Workflow_Context.json").write_text(
        json.dumps({"meta": state["meta"], "jira": state["jira"]}, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    for source in sorted(DATA_DIR.iterdir()):
        if source.is_file():
            shutil.copy2(source, ENV_DIR / source.name)


if __name__ == "__main__":
    main()
