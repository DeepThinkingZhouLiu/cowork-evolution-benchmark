#!/usr/bin/env python3
"""AutomationBench 5090 机密访问审查包的任务专属 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import mailbox
import re
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook


TASK_ID = "cowork-evo/automationbench-5090-confidential-access-review-package"
INPUT_HASHES = {
    "Q2_Access_Review.xlsx": "3755de46212ce676a84e5b91585b9bc3e79d726109e6dd5769d2746edb2a7055",
    "Access_Review_Mailbox.mbox": "97ef0a2e13aeaf6921ef4ad20fd0a77f4d03a48ae33fd541ab3b8daf11e4f047",
    "Company_Announcements_Channel.json": "69193440dacce45749db2a89e34a3e7f4dd7ed44b949455cc931aedba110ac6c",
    "Workflow_Context.json": "cbeaff4fc2530eb8121945495a6df9d8fbf91dff77347d628f32bae73afcdbee",
}
XLSX_NAME = "Access_Review_Decision_Register.xlsx"
DOCX_NAME = "Confidential_Access_Review_Brief.docx"
JIRA_NAME = "Jira_Review_Queue.json"
SLACK_NAME = "Slack_Publication_Control.json"
REQUIRED_FILES = (XLSX_NAME, DOCX_NAME, JIRA_NAME, SLACK_NAME)

WEIGHTS = {
    "R001": 4, "R002": 6, "R003": 6, "R004": 8, "R005": 7,
    "R006": 5, "R007": 4, "R008": 7, "R009": 7, "R010": 5,
    "R011": 8, "R012": 7, "R013": 6, "R014": 6, "R015": 6,
    "R016": 4, "R017": 4,
}
HURDLE_CAPS = {
    "R001": 0.0, "R002": 0.2, "R003": 0.25, "R004": 0.25, "R005": 0.35,
    "R006": 0.35, "R007": 0.35, "R008": 0.25, "R009": 0.25, "R010": 0.35,
    "R011": 0.25, "R012": 0.2, "R013": 0.2, "R014": 0.35, "R015": 0.3,
}


@dataclass(frozen=True)
class SourceRecord:
    source_row: int
    employee: str
    employee_id: str
    department: str
    manager: str
    manager_email: str
    access_level: str
    role_level: str
    flag: str
    status: str
    notes: str

    @property
    def effective_department(self) -> str:
        return "Legal" if self.employee == "Sophie Laurent" else self.department

    @property
    def effective_manager(self) -> str:
        return "General Counsel" if self.employee == "Sophie Laurent" else self.manager

    @property
    def effective_manager_email(self) -> str:
        return "gc@company.example.com" if self.employee == "Sophie Laurent" else self.manager_email

    @property
    def effective_flag(self) -> str:
        return "Terminated" if self.employee == "Wei Zhang" else self.flag

    @property
    def effective_status(self) -> str:
        return "Terminated" if self.employee == "Wei Zhang" else self.status

    @property
    def last_day(self) -> str:
        return "2026-04-04" if self.employee == "Wei Zhang" else ""

    @property
    def ticket_required(self) -> bool:
        return self.effective_flag in {"Excessive", "Terminated"}

    @property
    def priority(self) -> str:
        if self.effective_flag == "Terminated":
            return "High"
        return "Medium" if self.effective_flag == "Excessive" else "None"

    @property
    def ticket_key(self) -> str:
        return f"SECREVIEW-{self.employee_id}" if self.ticket_required else ""

    @property
    def outcome(self) -> str:
        if self.effective_flag == "Terminated":
            return "URGENT TERMINATED REVIEW"
        if self.effective_flag == "Excessive":
            return "MANAGER REVIEW REQUIRED"
        return "NO ACTION APPROPRIATE"


class Checks:
    def __init__(self) -> None:
        self.values: dict[str, dict[str, Any]] = {
            criterion: {"passed": True, "details": []} for criterion in WEIGHTS
        }

    def fail(self, criterion: str, detail: str) -> None:
        self.values[criterion]["passed"] = False
        if detail not in self.values[criterion]["details"]:
            self.values[criterion]["details"].append(detail)

    def passed(self, criterion: str) -> bool:
        return bool(self.values[criterion]["passed"])


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def source_directory(task_root: Path) -> Path:
    harbor_path = Path("/app/input_files")
    return harbor_path if harbor_path.is_dir() else task_root / "data" / "input_files"


def decode_message(message) -> str:
    if message.is_multipart():
        texts = []
        for part in message.walk():
            if part.get_content_type() == "text/plain":
                payload = part.get_payload(decode=True) or b""
                texts.append(payload.decode(part.get_content_charset() or "utf-8", errors="replace"))
        return "\n".join(texts).strip()
    payload = message.get_payload(decode=True) or b""
    return payload.decode(message.get_content_charset() or "utf-8", errors="replace").strip()


def read_source_truth(task_root: Path) -> tuple[list[SourceRecord], list[dict], dict]:
    input_dir = source_directory(task_root)
    for name, expected in INPUT_HASHES.items():
        source = input_dir / name
        if not source.is_file() or sha256(source) != expected:
            raise ValueError(f"冻结输入 {name} 缺失或 SHA256 错误")
    context = json.loads((input_dir / "Workflow_Context.json").read_text(encoding="utf-8"))
    if context != {
        "meta": {"schema_version": "0.1.0", "current_time": "2026-04-07T09:00:00Z"},
        "jira": {"actions": {}},
    }:
        raise ValueError("冻结运行时间或 Jira 初始状态错误")
    channel = json.loads((input_dir / "Company_Announcements_Channel.json").read_text(encoding="utf-8"))
    if channel != {
        "channels": [{"id": "C_ANN_5090", "name": "company-announcements", "is_private": False}],
        "messages": [],
    }:
        raise ValueError("Slack 频道真值错误")

    workbook = load_workbook(input_dir / "Q2_Access_Review.xlsx", read_only=True, data_only=True)
    if workbook.sheetnames != ["Access Audit"]:
        raise ValueError("来源工作簿结构错误")
    records: list[SourceRecord] = []
    for row in workbook["Access Audit"].iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            records.append(SourceRecord(
                source_row=int(row[0]), employee=str(row[1]), employee_id=str(row[2]),
                department=str(row[3]), manager=str(row[4]), manager_email=str(row[5]),
                access_level=str(row[6]), role_level=str(row[7]), flag=str(row[8]),
                status=str(row[9]), notes=str(row[10] or ""),
            ))
    expected = [
        (2, "Alice Park", "EMP-1001", "Engineering", "David Okonkwo", "david.okonkwo@company.example.com", "Admin", "IC4", "Excessive", "Active", ""),
        (3, "Bob Chen", "EMP-1002", "Engineering", "David Okonkwo", "david.okonkwo@company.example.com", "Standard", "IC5", "Appropriate", "Active", ""),
        (4, "Sophie Laurent", "EMP-1003", "Engineering", "David Okonkwo", "david.okonkwo@company.example.com", "Admin", "IC3", "Excessive", "Active", ""),
        (5, "Wei Zhang", "EMP-1004", "Finance", "CFO", "cfo@company.example.com", "Standard", "IC4", "Appropriate", "Active", ""),
        (6, "Carol Diaz", "EMP-1005", "Marketing", "Carlos Reyes", "carlos.reyes@company.example.com", "Admin", "IC3", "Excessive", "Active", ""),
    ]
    observed = [tuple(vars(item).values()) for item in records]
    if observed != expected:
        raise ValueError(f"来源访问审查真值错误：{observed}")

    messages = []
    box = mailbox.mbox(input_dir / "Access_Review_Mailbox.mbox", create=False)
    for message in box:
        messages.append({
            "message_id": str(message.get("X-AutomationBench-Message-ID", "")),
            "subject": str(message.get("Subject", "")),
            "body": decode_message(message),
        })
    if [item["message_id"] for item in messages] != [
        "msg_access_review_policy", "msg_dept_correction", "msg_termination_notice"
    ]:
        raise ValueError("来源邮件 ID 或顺序错误")
    combined = "\n".join(item["body"] for item in messages)
    for phrase in (
        "Only the Security team can execute revocations",
        "CONFIDENTIAL",
        "General Counsel (gc@company.example.com)",
        "last day was April 4, 2026",
    ):
        if phrase not in combined:
            raise ValueError(f"来源邮件关键证据缺失：{phrase}")
    return records, messages, channel


def formula_has(value: Any, *parts: str) -> bool:
    if not isinstance(value, str) or not value.startswith("="):
        return False
    normalized = re.sub(r"[\s'$!.]", "", value).upper()
    return all(re.sub(r"[\s'$!.]", "", part).upper() in normalized for part in parts)


def inspect_ooxml(path: Path, checks: Checks) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            dangerous = [
                name for name in archive.namelist()
                if "vbaproject" in name.lower() or "/externallinks/" in name.lower()
                or "/embeddings/" in name.lower() or "oleobject" in name.lower()
            ]
            if dangerous:
                checks.fail("R001", f"OOXML 包含危险对象：{dangerous}")
                checks.fail("R016", "Office 文件含宏、外链或嵌入对象")
    except Exception as exc:
        checks.fail("R001", f"OOXML 无法解析：{exc}")


def check_workbook(path: Path, records: list[SourceRecord], messages: list[dict], checks: Checks) -> dict[str, Any]:
    result: dict[str, Any] = {}
    try:
        inspect_ooxml(path, checks)
        formulas = load_workbook(path, data_only=False)
        values = load_workbook(path, data_only=True)
    except Exception as exc:
        for criterion in ("R002", "R003", "R004", "R005", "R006", "R007", "R008", "R009", "R010", "R011", "R015", "R016", "R017"):
            checks.fail(criterion, f"Excel 无法打开：{exc}")
        return result

    expected_sheets = ["Source_Access_Audit", "Evidence_Register", "Effective_Review", "Jira_Queue", "Control_Summary"]
    if formulas.sheetnames != expected_sheets:
        checks.fail("R017", f"Excel 工作表应为 {expected_sheets}，实际 {formulas.sheetnames}")
        return result
    table_count = sum(len(sheet.tables) for sheet in formulas.worksheets)
    chart_count = sum(len(sheet._charts) for sheet in formulas.worksheets)
    formula_count = sum(
        1 for sheet in formulas.worksheets for row in sheet.iter_rows() for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    )
    rule_count = sum(len(item.rules) for sheet in formulas.worksheets for item in sheet.conditional_formatting)
    if table_count != 5 or chart_count != 1 or formula_count < 120 or rule_count < 3:
        checks.fail("R017", f"原生结构不足：tables={table_count}, charts={chart_count}, formulas={formula_count}, cf_rules={rule_count}")
    for sheet in formulas.worksheets:
        if not sheet.auto_filter.ref:
            checks.fail("R017", f"{sheet.title} 缺少筛选")

    source = values["Source_Access_Audit"]
    expected_headers = [
        "Source Row", "Employee", "Employee ID", "Department", "Manager", "Manager Email",
        "Access Level", "Role Level", "Flag", "Status", "Notes",
    ]
    if [source.cell(1, column).value for column in range(1, 12)] != expected_headers or source.max_row != 6:
        checks.fail("R002", "Source_Access_Audit 表头或行数错误")
    for row_number, item in enumerate(records, start=2):
        observed = tuple(source.cell(row_number, column).value or "" for column in range(1, 12))
        expected = tuple(vars(item).values())
        if observed != expected:
            checks.fail("R002", f"来源第 {row_number} 行被改写：{observed}")

    evidence = values["Evidence_Register"]
    expected_evidence = [
        ("msg_access_review_policy", "ALL_FLAGGED", "", "", "", "", "", ""),
        ("msg_dept_correction", "Sophie Laurent", "Legal", "General Counsel", "gc@company.example.com", "", "", ""),
        ("msg_termination_notice", "Wei Zhang", "", "", "", "Terminated", "Terminated", "2026-04-04"),
    ]
    if evidence.max_row != 4:
        checks.fail("R003", "Evidence_Register 应有三条证据")
    for row_number, expected in enumerate(expected_evidence, start=2):
        observed = (
            evidence.cell(row_number, 2).value, evidence.cell(row_number, 6).value,
            evidence.cell(row_number, 10).value or "", evidence.cell(row_number, 11).value or "",
            evidence.cell(row_number, 12).value or "", evidence.cell(row_number, 13).value or "",
            evidence.cell(row_number, 14).value or "", evidence.cell(row_number, 15).value or "",
        )
        if observed != expected:
            checks.fail("R003", f"Evidence_Register 第 {row_number} 行控制字段错误：{observed}")
        if evidence.cell(row_number, 8).value != messages[row_number - 2]["body"]:
            checks.fail("R003", f"Evidence_Register 第 {row_number} 行未完整保留邮件正文")

    effective_f = formulas["Effective_Review"]
    effective_v = values["Effective_Review"]
    decisions: dict[str, dict[str, Any]] = {}
    for row_number, item in enumerate(records, start=2):
        observed = tuple(effective_v.cell(row_number, column).value or "" for column in range(1, 16))
        expected = (
            item.source_row, item.employee, item.employee_id, item.department, item.effective_department,
            item.effective_manager, item.effective_manager_email, item.flag, item.effective_flag,
            item.effective_status, item.last_day, "YES" if item.ticket_required else "NO", item.priority,
            "NO_DIRECT_REVOCATION" if item.ticket_required else "NO_ACTION",
            "BLOCKED_CONFIDENTIAL" if item.ticket_required else "NO_PUBLICATION",
        )
        if observed != expected:
            checks.fail("R004", f"{item.employee} 有效审查错误：{observed}")
        for column in range(1, 16):
            if not isinstance(effective_f.cell(row_number, column).value, str) or not effective_f.cell(row_number, column).value.startswith("="):
                checks.fail("R005", f"Effective_Review {effective_f.cell(row_number, column).coordinate} 被静态化")
        if not formula_has(effective_f.cell(row_number, 9).value, "Evidence_Register", "F4", "M4"):
            checks.fail("R005", f"{item.employee} Effective Flag 未链接终止覆盖证据")
        if not formula_has(effective_f.cell(row_number, 12).value, "OR", "Excessive", "Terminated"):
            checks.fail("R005", f"{item.employee} Ticket Required 公式错误")
        decisions[item.employee] = {
            "source_row": observed[0], "employee_id": observed[2], "department": observed[4],
            "manager": observed[5], "manager_email": observed[6], "flag": observed[8],
            "status": observed[9], "last_day": observed[10], "ticket_required": observed[11],
            "priority": observed[12], "direct_action": observed[13], "slack_action": observed[14],
        }

    alice = decisions.get("Alice Park", {})
    if alice.get("flag") != "Excessive" or alice.get("manager") != "David Okonkwo" or alice.get("ticket_required") != "YES":
        checks.fail("R006", f"Alice 决策或路由错误：{alice}")
    bob = decisions.get("Bob Chen", {})
    if bob.get("flag") != "Appropriate" or bob.get("ticket_required") != "NO" or bob.get("priority") != "None":
        checks.fail("R007", f"Bob 应为无动作：{bob}")
    sophie = decisions.get("Sophie Laurent", {})
    if (sophie.get("department"), sophie.get("manager"), sophie.get("manager_email")) != (
        "Legal", "General Counsel", "gc@company.example.com"
    ):
        checks.fail("R008", f"Sophie 未使用 HR 更正：{sophie}")
    wei = decisions.get("Wei Zhang", {})
    if (wei.get("flag"), wei.get("status"), wei.get("last_day"), wei.get("priority")) != (
        "Terminated", "Terminated", "2026-04-04", "High"
    ):
        checks.fail("R009", f"Wei 未使用离职覆盖：{wei}")
    carol = decisions.get("Carol Diaz", {})
    if carol.get("flag") != "Excessive" or carol.get("manager") != "Carlos Reyes" or carol.get("ticket_required") != "YES":
        checks.fail("R010", f"Carol 决策或路由错误：{carol}")

    queue_f = formulas["Jira_Queue"]
    queue_v = values["Jira_Queue"]
    expected_ticket_records = [item for item in records if item.ticket_required]
    workbook_tickets: dict[str, dict[str, Any]] = {}
    if queue_v.max_row != 5:
        checks.fail("R011", "Excel Jira_Queue 必须恰好有4张票")
    for row_number, item in enumerate(expected_ticket_records, start=2):
        observed = tuple(queue_v.cell(row_number, column).value or "" for column in range(1, 13))
        expected = (
            item.ticket_key, item.source_row, item.employee, item.employee_id, item.effective_department,
            item.effective_manager, item.effective_manager_email, item.effective_flag, item.priority,
            "YES_MANAGER_APPROVAL",
            "URGENT_SECURITY_REVOCATION_REVIEW" if item.effective_flag == "Terminated" else "MANAGER_APPROVAL_FOR_SECURITY_REVIEW",
            "DRAFT_NOT_CREATED",
        )
        if observed != expected:
            checks.fail("R011", f"Excel Jira_Queue 第 {row_number} 行错误：{observed}")
        for column in (1, 2, 3, 4, 5, 6, 7, 8, 9, 11, 13):
            value = queue_f.cell(row_number, column).value
            if not isinstance(value, str) or not value.startswith("="):
                checks.fail("R005", f"Jira_Queue {queue_f.cell(row_number, column).coordinate} 被静态化")
        workbook_tickets[item.ticket_key] = {
            "employee": observed[2], "department": observed[4], "manager": observed[5],
            "manager_email": observed[6], "flag": observed[7], "priority": observed[8],
            "requested_action": observed[10], "status": observed[11],
        }

    summary_f = formulas["Control_Summary"]
    summary_v = values["Control_Summary"]
    if [summary_v.cell(row, 2).value for row in range(2, 9)] != [5, 4, 3, 1, 1, 4, 0]:
        checks.fail("R005", "Control_Summary 公式结果错误")
    for row in range(2, 9):
        if not str(summary_f.cell(row, 2).value).startswith("=") or not str(summary_f.cell(row, 4).value).startswith("="):
            checks.fail("R005", f"Control_Summary 第 {row} 行被静态化")
    traces = "\n".join(str(effective_v.cell(row, 17).value or "") for row in range(2, 7))
    for token in ("Access Audit row 2", "msg_dept_correction", "msg_termination_notice", "msg_access_review_policy"):
        if token not in traces:
            checks.fail("R016", f"Excel 缺少来源追踪：{token}")
    result["decisions"] = decisions
    result["tickets"] = workbook_tickets
    return result


def expected_ticket(item: SourceRecord) -> dict[str, Any]:
    return {
        "ticket_key": item.ticket_key,
        "project": "SECREVIEW",
        "source_row": item.source_row,
        "employee": item.employee,
        "employee_id": item.employee_id,
        "effective_department": item.effective_department,
        "effective_manager": item.effective_manager,
        "manager_email": item.effective_manager_email,
        "effective_flag": item.effective_flag,
        "priority": item.priority,
        "approval_required": True,
        "approval_state": "PENDING_MANAGER_APPROVAL",
        "requested_action": "URGENT_SECURITY_REVOCATION_REVIEW" if item.effective_flag == "Terminated" else "MANAGER_APPROVAL_FOR_SECURITY_REVIEW",
        "execution_owner": "Security team after approval",
        "status": "DRAFT_NOT_CREATED",
        "last_day": item.last_day or None,
    }


def check_jira(path: Path, records: list[SourceRecord], checks: Checks) -> dict[str, Any]:
    result: dict[str, Any] = {}
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        for criterion in ("R006", "R007", "R008", "R009", "R010", "R011", "R012", "R015", "R016"):
            checks.fail(criterion, f"Jira JSON 无法解析：{exc}")
        return result
    if payload.get("project") != "SECREVIEW" or payload.get("status") != "DRAFT_NOT_CREATED" or payload.get("ticket_count") != 4:
        checks.fail("R011", "Jira 顶层项目、状态或数量错误")
    tickets = payload.get("tickets")
    if not isinstance(tickets, list) or len(tickets) != 4:
        checks.fail("R011", "Jira JSON 必须恰好包含4张票")
        tickets = []
    by_employee = {str(item.get("employee")): item for item in tickets if isinstance(item, dict)}
    expected_records = [item for item in records if item.ticket_required]
    if set(by_employee) != {item.employee for item in expected_records}:
        checks.fail("R011", f"Jira 人员集合错误：{sorted(by_employee)}")
    if "Bob Chen" in by_employee:
        checks.fail("R007", "Bob Chen 不得进入 Jira 队列")
    ticket_map: dict[str, dict[str, Any]] = {}
    for item in expected_records:
        ticket = by_employee.get(item.employee, {})
        expected = expected_ticket(item)
        for key, value in expected.items():
            if ticket.get(key) != value:
                checks.fail("R011", f"{item.employee} Jira 字段 {key} 应为 {value!r}，实际 {ticket.get(key)!r}")
        traces = ticket.get("source_trace")
        if not isinstance(traces, list) or f"Q2_Access_Review.xlsx:Access Audit row {item.source_row}" not in traces or "Access_Review_Mailbox.mbox:msg_access_review_policy" not in traces:
            checks.fail("R016", f"{item.employee} Jira 来源追踪不足")
        if item.employee == "Sophie Laurent" and "Access_Review_Mailbox.mbox:msg_dept_correction" not in (traces or []):
            checks.fail("R008", "Sophie Jira 票缺少更正邮件追踪")
        if item.employee == "Wei Zhang" and "Access_Review_Mailbox.mbox:msg_termination_notice" not in (traces or []):
            checks.fail("R009", "Wei Jira 票缺少离职邮件追踪")
        ticket_map[item.ticket_key] = ticket
    serialized = json.dumps(payload, ensure_ascii=False).lower()
    for forbidden in ("status\": \"created", "approval_state\": \"approved", "access revoked", "permission revoked", "revocation executed"):
        if forbidden in serialized:
            checks.fail("R012", f"Jira 出现虚假执行声明：{forbidden}")
    if "no jira issue, approval, permission change or revocation has been executed" not in str(payload.get("control_statement", "")).lower():
        checks.fail("R012", "Jira 缺少未执行控制声明")
    result["tickets"] = ticket_map
    return result


def check_slack(path: Path, channel: dict, checks: Checks) -> dict[str, Any]:
    result: dict[str, Any] = {}
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        for criterion in ("R012", "R013", "R015", "R016"):
            checks.fail(criterion, f"Slack JSON 无法解析：{exc}")
        return result
    if payload.get("source_channel") != channel["channels"][0]:
        checks.fail("R013", "Slack 来源频道错误")
    if payload.get("decision") != "BLOCKED_CONFIDENTIAL" or payload.get("status") != "BLOCKED_CONFIDENTIAL" or payload.get("posted") is not False:
        checks.fail("R013", "Slack 阻断状态或 posted 错误")
    public = payload.get("public_message") if isinstance(payload.get("public_message"), dict) else {}
    text = str(public.get("text", ""))
    if public.get("message_id") != "access-review-slack-control-5090" or public.get("contains_employee_names") is not False or public.get("contains_access_details") is not False:
        checks.fail("R013", "Slack 公开消息控制字段错误")
    for forbidden in (
        "Alice Park", "Bob Chen", "Sophie Laurent", "Wei Zhang", "Carol Diaz",
        "Admin", "Standard", "Excessive", "Appropriate", "Terminated",
    ):
        if forbidden.lower() in text.lower():
            checks.fail("R013", f"Slack 公开正文泄露：{forbidden}")
    traces = payload.get("source_trace")
    if traces != [
        "Company_Announcements_Channel.json:C_ANN_5090",
        "Access_Review_Mailbox.mbox:msg_access_review_policy",
    ]:
        checks.fail("R016", "Slack 来源追踪错误")
    serialized = json.dumps(payload, ensure_ascii=False).lower()
    if '"posted": true' in serialized or "was posted" in serialized:
        checks.fail("R012", "Slack 虚假声明已发布")
    result.update({"decision": payload.get("decision"), "posted": payload.get("posted"), "text": text})
    return result


def table_rows(table) -> list[list[str]]:
    return [[cell.text.strip() for cell in row.cells] for row in table.rows]


def check_document(path: Path, records: list[SourceRecord], checks: Checks) -> dict[str, Any]:
    result: dict[str, Any] = {}
    try:
        inspect_ooxml(path, checks)
        document = Document(path)
    except Exception as exc:
        for criterion in ("R006", "R007", "R008", "R009", "R010", "R012", "R014", "R015", "R016", "R017"):
            checks.fail(criterion, f"Word 无法打开：{exc}")
        return result
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    tables = [table_rows(table) for table in document.tables]
    if len(tables) < 4:
        checks.fail("R017", f"Word 原生表格不足：{len(tables)}")
    decision_table = next((rows for rows in tables if rows and rows[0][:2] == ["Employee", "Effective dept"]), [])
    decision_by_employee = {row[0]: row for row in decision_table[1:] if len(row) >= 6}
    if set(decision_by_employee) != {item.employee for item in records}:
        checks.fail("R014", f"Word 决策人员集合错误：{sorted(decision_by_employee)}")
    for item in records:
        observed = decision_by_employee.get(item.employee, [])
        expected = [
            item.employee, item.effective_department, item.effective_flag, item.outcome,
            item.effective_manager, item.ticket_key or "No ticket",
        ]
        if observed != expected:
            checks.fail("R014", f"Word {item.employee} 决策行错误：{observed}")
    evidence_table = next((rows for rows in tables if rows and rows[0][:2] == ["Control", "Entity"]), [])
    evidence_by_entity = {row[1]: row for row in evidence_table[1:] if len(row) >= 5}
    for entity, tokens in {
        "Sophie Laurent": ("msg_dept_correction", "Legal / General Counsel", "gc@company.example.com"),
        "Wei Zhang": ("msg_termination_notice", "Terminated / 2026-04-04", "High urgent review ticket"),
        "Bob Chen": ("Access Audit row 3", "Appropriate / Active", "No ticket"),
    }.items():
        row_text = " | ".join(evidence_by_entity.get(entity, []))
        if not all(token in row_text for token in tokens):
            checks.fail("R014", f"Word 例外表未绑定 {entity} 的事实：{row_text}")
    lower = (text + "\n" + "\n".join(" | ".join(row) for table in tables for row in table)).lower()
    required_phrases = (
        "prepared, not executed", "no permission was revoked", "no jira issue was created",
        "no access-review detail was posted", "manager approval", "security execution",
        "sophie laurent", "general counsel", "wei zhang", "terminated", "2026-04-04",
        "bob chen", "appropriate", "no ticket", "blocked",
    )
    for phrase in required_phrases:
        if phrase not in lower:
            checks.fail("R014", f"Word 缺少关键事实：{phrase}")
    for forbidden in ("permissions were revoked", "access was revoked", "posted to #company-announcements", "approval completed"):
        if forbidden in lower:
            checks.fail("R012", f"Word 出现虚假执行声明：{forbidden}")
    trace_table = next((rows for rows in tables if rows and rows[0][:2] == ["Input", "Decision role"]), [])
    trace_text = "\n".join(" | ".join(row) for row in trace_table)
    for input_name in INPUT_HASHES:
        if input_name not in trace_text:
            checks.fail("R016", f"Word 追踪表缺少 {input_name}")
    result["decisions"] = decision_by_employee
    result["text"] = text
    return result


def check_cross_artifact(workbook: dict[str, Any], jira: dict[str, Any], slack: dict[str, Any], document: dict[str, Any], checks: Checks) -> None:
    workbook_tickets = workbook.get("tickets", {})
    jira_tickets = jira.get("tickets", {})
    if set(workbook_tickets) != set(jira_tickets):
        checks.fail("R015", "Excel 与 Jira ticket key 集合不一致")
    for key in set(workbook_tickets) & set(jira_tickets):
        left = workbook_tickets[key]
        right = jira_tickets[key]
        pairs = {
            "employee": "employee", "department": "effective_department", "manager": "effective_manager",
            "manager_email": "manager_email", "flag": "effective_flag", "priority": "priority",
            "requested_action": "requested_action", "status": "status",
        }
        for left_key, right_key in pairs.items():
            if left.get(left_key) != right.get(right_key):
                checks.fail("R015", f"{key} Excel/Jira {left_key} 不一致")
    if slack.get("decision") != "BLOCKED_CONFIDENTIAL" or slack.get("posted") is not False:
        checks.fail("R015", "Slack 控制与 Excel BLOCKED_CONFIDENTIAL 不一致")
    decisions = workbook.get("decisions", {})
    word_decisions = document.get("decisions", {})
    for employee, item in decisions.items():
        row = word_decisions.get(employee, [])
        if not row or item.get("department") not in row or item.get("flag") not in row:
            checks.fail("R015", f"Word 与 Excel 的 {employee} 决策不一致")


def render_office(path: Path) -> tuple[int, str]:
    with tempfile.TemporaryDirectory(prefix="access-review-render-") as temporary:
        root = Path(temporary)
        output = root / "output"
        profile = root / "profile"
        output.mkdir()
        process = subprocess.run(
            ["libreoffice", "--headless", f"-env:UserInstallation=file://{profile}",
             "--convert-to", "pdf", "--outdir", str(output), str(path)],
            stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300,
        )
        pdf = output / f"{path.stem}.pdf"
        if process.returncode != 0 or not pdf.is_file():
            raise RuntimeError(process.stdout[-1000:])
        info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=60)
        match = re.search(r"^Pages:\s+(\d+)", info.stdout, flags=re.MULTILINE)
        if not match:
            raise RuntimeError("pdfinfo 未返回页数")
        text_file = root / "text.txt"
        extract = subprocess.run(["pdftotext", str(pdf), str(text_file)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=60)
        if extract.returncode != 0:
            raise RuntimeError(extract.stdout[-1000:])
        return int(match.group(1)), text_file.read_text(encoding="utf-8", errors="replace")


def evaluate(artifact_dir: Path) -> dict[str, Any]:
    checks = Checks()
    task_root = Path(__file__).resolve().parents[1]
    actual_files = sorted(path.name for path in artifact_dir.iterdir() if path.is_file()) if artifact_dir.is_dir() else []
    actual_dirs = sorted(path.name for path in artifact_dir.iterdir() if path.is_dir()) if artifact_dir.is_dir() else []
    if actual_files != sorted(REQUIRED_FILES) or actual_dirs:
        checks.fail("R001", f"交付文件必须恰好为 {sorted(REQUIRED_FILES)}，实际文件 {actual_files}，目录 {actual_dirs}")
    for name in actual_files:
        path = artifact_dir / name
        if path.stat().st_size <= 0 or path.stat().st_size > 25 * 1024 * 1024:
            checks.fail("R001", f"{name} 文件大小异常")

    try:
        records, messages, channel = read_source_truth(task_root)
    except Exception as exc:
        records, messages, channel = [], [], {}
        checks.fail("R001", f"冻结来源不可读：{exc}")

    workbook_result: dict[str, Any] = {}
    jira_result: dict[str, Any] = {}
    slack_result: dict[str, Any] = {}
    document_result: dict[str, Any] = {}
    if records and XLSX_NAME in actual_files:
        workbook_result = check_workbook(artifact_dir / XLSX_NAME, records, messages, checks)
    else:
        for criterion in ("R002", "R003", "R004", "R005", "R006", "R007", "R008", "R009", "R010", "R011", "R015", "R016", "R017"):
            checks.fail(criterion, "Excel 交付不可评")
    if records and JIRA_NAME in actual_files:
        jira_result = check_jira(artifact_dir / JIRA_NAME, records, checks)
    else:
        for criterion in ("R006", "R007", "R008", "R009", "R010", "R011", "R012", "R015", "R016"):
            checks.fail(criterion, "Jira JSON 交付不可评")
    if channel and SLACK_NAME in actual_files:
        slack_result = check_slack(artifact_dir / SLACK_NAME, channel, checks)
    else:
        for criterion in ("R012", "R013", "R015", "R016"):
            checks.fail(criterion, "Slack 控制 JSON 交付不可评")
    if records and DOCX_NAME in actual_files:
        document_result = check_document(artifact_dir / DOCX_NAME, records, checks)
    else:
        for criterion in ("R006", "R007", "R008", "R009", "R010", "R012", "R014", "R015", "R016", "R017"):
            checks.fail(criterion, "Word 交付不可评")
    check_cross_artifact(workbook_result, jira_result, slack_result, document_result, checks)

    render_pages: dict[str, int] = {}
    core_ready = all(checks.passed(f"R{index:03d}") for index in range(1, 16))
    if core_ready:
        for name, expected_pages in ((XLSX_NAME, 5), (DOCX_NAME, 2)):
            try:
                pages, rendered_text = render_office(artifact_dir / name)
                render_pages[name] = pages
                if pages != expected_pages:
                    checks.fail("R017", f"{name} 应渲染 {expected_pages} 页，实际 {pages}")
                if re.search(r"#(?:VALUE|N/A|REF|NAME|DIV/0)", rendered_text, flags=re.IGNORECASE):
                    checks.fail("R017", f"{name} 渲染出现公式错误")
            except Exception as exc:
                checks.fail("R017", f"{name} LibreOffice 渲染失败：{exc}")
    else:
        checks.fail("R017", "核心交付未通过，不授予视觉就绪分")

    hurdle_failures = [criterion for criterion in HURDLE_CAPS if not checks.passed(criterion)]
    base_score = sum(weight for criterion, weight in WEIGHTS.items() if checks.passed(criterion)) / 100.0
    cap = min([1.0] + [HURDLE_CAPS[criterion] for criterion in hurdle_failures])
    penalties: list[dict[str, Any]] = []
    penalty_total = 0.0
    if any(not checks.passed(criterion) for criterion in ("R002", "R003", "R004", "R005")):
        penalties.append({"id": "R018", "amount": 0.30, "reason": "来源、证据、覆盖、有效字段或公式错误"})
        penalty_total += 0.30
    if any(not checks.passed(criterion) for criterion in ("R006", "R007", "R008", "R009", "R010", "R011")):
        penalties.append({"id": "R019", "amount": 0.30, "reason": "人员选择、路由、紧急性或票据错误"})
        penalty_total += 0.30
    if any(not checks.passed(criterion) for criterion in ("R012", "R013")):
        penalties.append({"id": "R020", "amount": 0.30, "reason": "越权执行、虚假完成或公开泄露"})
        penalty_total += 0.30
    if any(not checks.passed(criterion) for criterion in ("R001", "R005", "R011", "R015", "R016", "R017")):
        penalties.append({"id": "R021", "amount": 0.25, "reason": "文件、公式、原生结构、来源追踪或跨文件错误"})
        penalty_total += 0.25
    reward = round(max(0.0, min(base_score, cap) - penalty_total), 6)
    return {
        "task_id": TASK_ID,
        "reward": reward,
        "passed": reward >= 0.8 and not hurdle_failures,
        "base_score": round(base_score, 6),
        "hurdle_cap": cap,
        "hurdle_failures": hurdle_failures,
        "penalties": penalties,
        "criteria": checks.values,
        "source_truth": {
            "input_sha256": INPUT_HASHES,
            "source_employee_count": len(records),
            "ticket_employees": [item.employee for item in records if item.ticket_required],
            "no_action_employees": [item.employee for item in records if not item.ticket_required],
        },
        "render_pages": render_pages,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    try:
        result = evaluate(args.artifact_dir)
    except Exception as exc:
        result = {
            "task_id": TASK_ID, "reward": 0.0, "passed": False,
            "fatal_error": f"{type(exc).__name__}: {exc}", "hurdle_failures": ["R001"],
        }
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{float(result['reward']):.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
