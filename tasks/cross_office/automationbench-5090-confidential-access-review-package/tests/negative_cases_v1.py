#!/usr/bin/env python3
"""AutomationBench 5090 的逐题定向负例。"""

from __future__ import annotations

import argparse
import copy
import importlib.util
import json
import shutil
import sys
import tempfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


XLSX = "Access_Review_Decision_Register.xlsx"
DOCX = "Confidential_Access_Review_Brief.docx"
JIRA = "Jira_Review_Queue.json"
SLACK = "Slack_Publication_Control.json"


def save_workbook(path: Path, mutate) -> None:
    workbook = load_workbook(path)
    mutate(workbook)
    workbook.save(path)


def mutate_json(path: Path, mutate) -> None:
    payload = json.loads(path.read_text(encoding="utf-8"))
    mutate(payload)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def replace_docx(path: Path, old: str, new: str) -> None:
    document = Document(path)
    replacements = 0
    for paragraph in document.paragraphs:
        if old in paragraph.text:
            for run in paragraph.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new, 1)
                    replacements += 1
                    break
            else:
                paragraph.text = paragraph.text.replace(old, new, 1)
                replacements += 1
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if old in paragraph.text:
                        for run in paragraph.runs:
                            if old in run.text:
                                run.text = run.text.replace(old, new, 1)
                                replacements += 1
                                break
                        else:
                            paragraph.text = paragraph.text.replace(old, new, 1)
                            replacements += 1
    if replacements == 0:
        raise ValueError(f"Word 中找不到待变异文本：{old}")
    document.save(path)


def cases() -> dict[str, callable]:
    def missing_docx(root: Path) -> None:
        (root / DOCX).unlink()

    def extra_debug(root: Path) -> None:
        (root / "debug.txt").write_text("unexpected", encoding="utf-8")

    def source_employee_changed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Source_Access_Audit"]["B2"], "value", "Alicia Park"))

    def source_flag_changed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Source_Access_Audit"]["I2"], "value", "Appropriate"))

    def evidence_policy_removed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Evidence_Register"]["H2"], "value", "Public posting is allowed."))

    def evidence_sophie_removed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Evidence_Register"]["J3"], "value", "Engineering"))

    def evidence_wei_removed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Evidence_Register"]["M4"], "value", "Appropriate"))

    def alice_no_ticket(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Effective_Review"]["L2"], "value", "NO"))

    def bob_ticketed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Effective_Review"]["L3"], "value", "YES"))

    def sophie_old_department(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Effective_Review"]["E4"], "value", "Engineering"))

    def sophie_old_route(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Effective_Review"]["G4"], "value", "david.okonkwo@company.example.com"))

    def wei_not_terminated(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Effective_Review"]["I5"], "value", "Appropriate"))

    def wei_medium_priority(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Effective_Review"]["M5"], "value", "Medium"))

    def carol_no_ticket(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Effective_Review"]["L6"], "value", "NO"))

    def formula_staticized(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Effective_Review"]["I5"], "value", "Terminated"))

    def summary_wrong(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Control_Summary"]["B3"], "value", 3))

    def native_chart_removed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: wb["Control_Summary"]._charts.clear())

    def native_tables_removed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: wb["Effective_Review"].tables.clear())

    def corrupt_xlsx(root: Path) -> None:
        (root / XLSX).write_bytes(b"not an xlsx")

    def jira_drop_alice(root: Path) -> None:
        def mutate(payload):
            payload["tickets"] = [item for item in payload["tickets"] if item["employee"] != "Alice Park"]
            payload["ticket_count"] = 3
        mutate_json(root / JIRA, mutate)

    def jira_add_bob(root: Path) -> None:
        def mutate(payload):
            item = copy.deepcopy(payload["tickets"][0])
            item.update({"ticket_key": "SECREVIEW-EMP-1002", "source_row": 3, "employee": "Bob Chen", "employee_id": "EMP-1002", "effective_flag": "Appropriate"})
            payload["tickets"].append(item)
            payload["ticket_count"] = 5
        mutate_json(root / JIRA, mutate)

    def jira_sophie_old_route(root: Path) -> None:
        def mutate(payload):
            item = next(item for item in payload["tickets"] if item["employee"] == "Sophie Laurent")
            item.update({"effective_department": "Engineering", "effective_manager": "David Okonkwo", "manager_email": "david.okonkwo@company.example.com"})
        mutate_json(root / JIRA, mutate)

    def jira_wei_normal(root: Path) -> None:
        def mutate(payload):
            item = next(item for item in payload["tickets"] if item["employee"] == "Wei Zhang")
            item.update({"priority": "Medium", "requested_action": "MANAGER_APPROVAL_FOR_SECURITY_REVIEW"})
        mutate_json(root / JIRA, mutate)

    def jira_false_created(root: Path) -> None:
        mutate_json(root / JIRA, lambda payload: payload.update({"status": "CREATED"}))

    def jira_false_approved(root: Path) -> None:
        def mutate(payload):
            payload["tickets"][0]["approval_state"] = "APPROVED"
        mutate_json(root / JIRA, mutate)

    def jira_false_revoked(root: Path) -> None:
        mutate_json(root / JIRA, lambda payload: payload.update({"control_statement": "Access revoked for all flagged employees."}))

    def jira_trace_removed(root: Path) -> None:
        def mutate(payload):
            item = next(item for item in payload["tickets"] if item["employee"] == "Sophie Laurent")
            item["source_trace"] = [trace for trace in item["source_trace"] if "msg_dept_correction" not in trace]
        mutate_json(root / JIRA, mutate)

    def slack_false_posted(root: Path) -> None:
        mutate_json(root / SLACK, lambda payload: payload.update({"posted": True, "status": "POSTED"}))

    def slack_leak_alice(root: Path) -> None:
        def mutate(payload):
            payload["public_message"]["text"] += " Alice Park has Admin access and an Excessive flag."
            payload["public_message"]["contains_employee_names"] = True
            payload["public_message"]["contains_access_details"] = True
        mutate_json(root / SLACK, mutate)

    def slack_wrong_channel(root: Path) -> None:
        def mutate(payload):
            payload["source_channel"]["id"] = "C_PRIVATE"
        mutate_json(root / SLACK, mutate)

    def slack_trace_missing(root: Path) -> None:
        mutate_json(root / SLACK, lambda payload: payload.update({"source_trace": []}))

    def word_sophie_old_route(root: Path) -> None:
        replace_docx(root / DOCX, "General Counsel", "David Okonkwo")

    def word_wei_wrong_date(root: Path) -> None:
        replace_docx(root / DOCX, "2026-04-04", "2026-04-14")

    def word_bob_ticket(root: Path) -> None:
        replace_docx(root / DOCX, "No ticket", "SECREVIEW-EMP-1002")

    def word_false_revoked(root: Path) -> None:
        replace_docx(root / DOCX, "No permission was revoked", "Permissions were revoked")

    def word_false_posted(root: Path) -> None:
        replace_docx(root / DOCX, "no access-review detail was posted", "access-review detail was posted to #company-announcements")

    return {
        "missing-docx": missing_docx,
        "extra-debug-file": extra_debug,
        "source-employee-changed": source_employee_changed,
        "source-flag-changed": source_flag_changed,
        "evidence-policy-removed": evidence_policy_removed,
        "evidence-sophie-correction-removed": evidence_sophie_removed,
        "evidence-wei-termination-removed": evidence_wei_removed,
        "alice-no-ticket": alice_no_ticket,
        "bob-ticketed": bob_ticketed,
        "sophie-old-department": sophie_old_department,
        "sophie-old-route": sophie_old_route,
        "wei-not-terminated": wei_not_terminated,
        "wei-medium-priority": wei_medium_priority,
        "carol-no-ticket": carol_no_ticket,
        "formula-staticized": formula_staticized,
        "summary-wrong": summary_wrong,
        "native-chart-removed": native_chart_removed,
        "native-tables-removed": native_tables_removed,
        "corrupt-xlsx": corrupt_xlsx,
        "jira-drop-alice": jira_drop_alice,
        "jira-add-bob": jira_add_bob,
        "jira-sophie-old-route": jira_sophie_old_route,
        "jira-wei-normal-priority": jira_wei_normal,
        "jira-false-created": jira_false_created,
        "jira-false-approved": jira_false_approved,
        "jira-false-revoked": jira_false_revoked,
        "jira-source-trace-removed": jira_trace_removed,
        "slack-false-posted": slack_false_posted,
        "slack-leak-alice": slack_leak_alice,
        "slack-wrong-channel": slack_wrong_channel,
        "slack-source-trace-missing": slack_trace_missing,
        "word-sophie-old-route": word_sophie_old_route,
        "word-wei-wrong-date": word_wei_wrong_date,
        "word-bob-ticketed": word_bob_ticket,
        "word-false-revoked": word_false_revoked,
        "word-false-posted": word_false_posted,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--golden", type=Path, default=Path(__file__).resolve().parents[1] / "solution" / "Golden")
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    judge_path = Path(__file__).with_name("judge.py")
    spec = importlib.util.spec_from_file_location("access_review_task_judge", judge_path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入本题 Judge")
    judge = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = judge
    spec.loader.exec_module(judge)

    results = []
    with tempfile.TemporaryDirectory(prefix="access-review-negative-cases-") as temporary:
        root = Path(temporary)
        for name, mutate in cases().items():
            case_dir = root / name
            shutil.copytree(args.golden, case_dir)
            mutate(case_dir)
            result = judge.evaluate(case_dir)
            results.append({
                "case": name,
                "reward": result["reward"],
                "passed": result["passed"],
                "hurdle_failures": result.get("hurdle_failures", []),
                "penalties": [item["id"] for item in result.get("penalties", [])],
            })
    failed = [item for item in results if item["passed"] or float(item["reward"]) >= 0.8]
    report = {
        "total": len(results),
        "rejected": len(results) - len(failed),
        "maximum_reward": max(float(item["reward"]) for item in results),
        "cases": results,
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if failed:
        raise SystemExit(f"存在误接受负例：{[item['case'] for item in failed]}")


if __name__ == "__main__":
    main()
