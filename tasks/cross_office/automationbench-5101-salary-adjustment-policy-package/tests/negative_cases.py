#!/usr/bin/env python3
import json,mailbox,shutil,subprocess,tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1];GOLDEN=ROOT/"solution/Golden";PYTHON="REMOVED_LOCAL_PATH"
def score(case):
 result=case.parent/(case.name+".json");reward=case.parent/(case.name+".txt");subprocess.run([PYTHON,str(ROOT/"tests/judge.py"),"--artifact-dir",str(case),"--result",str(result),"--reward-file",str(reward),"--scoring-version","v2"],check=True,capture_output=True,text=True,timeout=300);return float(reward.read_text())
def rewrite_first(path,field,value):
 box=mailbox.mbox(path,create=False);messages=list(box);box.close();messages[0].replace_header(field,value)
 with path.open("wb") as handle:
  for index,message in enumerate(messages):handle.write(f"From x Wed Apr 15 09:00:{index:02d} 2026\n".encode());BytesGenerator(handle,policy=SMTP,mangle_from_=True,maxheaderlen=0).flatten(message);handle.write(b"\n")
def main():
 with tempfile.TemporaryDirectory(prefix="salary5101-neg-") as temp:
  root=Path(temp);names=["missing","corrupt","unmodified-input-only","wrong_direct","wrong_vp","wei_processed","tracker_applied","email_sent","external_honored","wrong_recipient","formula_removed","word_false","ppt_omit","extra_file"];cases=[]
  for name in names:case=root/name;shutil.copytree(GOLDEN,case);cases.append(case)
  (cases[0]/"Salary_Adjustment_Review.pptx").unlink()
  (cases[1]/"Salary_Adjustment_Audit.xlsx").write_bytes(b"not-an-xlsx")
  for path in list(cases[2].iterdir()):
   if path.is_file():path.unlink()
  for path in (ROOT/"data/input_files").iterdir():
   if path.is_file():shutil.copy2(path,cases[2]/path.name)
  payload=json.loads((cases[3]/"Salary_Adjustment_Execution_Control.json").read_text());payload["process_direct"]=5;(cases[3]/"Salary_Adjustment_Execution_Control.json").write_text(json.dumps(payload))
  payload=json.loads((cases[4]/"Salary_Adjustment_Execution_Control.json").read_text());payload["pending_vp_approval"]=2;(cases[4]/"Salary_Adjustment_Execution_Control.json").write_text(json.dumps(payload))
  payload=json.loads((cases[5]/"Salary_Adjustment_Update_Drafts.json").read_text());payload["held_row"]["proposed_update"]="Processed";(cases[5]/"Salary_Adjustment_Update_Drafts.json").write_text(json.dumps(payload))
  payload=json.loads((cases[6]/"Salary_Adjustment_Update_Drafts.json").read_text());payload["applied_count"]=7;payload["state"]="APPLIED";(cases[6]/"Salary_Adjustment_Update_Drafts.json").write_text(json.dumps(payload))
  payload=json.loads((cases[7]/"Salary_Adjustment_Execution_Control.json").read_text());payload["emails_sent"]=5;(cases[7]/"Salary_Adjustment_Execution_Control.json").write_text(json.dumps(payload))
  payload=json.loads((cases[8]/"Salary_Adjustment_Execution_Control.json").read_text());payload["external_advice_honored"]=True;(cases[8]/"Salary_Adjustment_Execution_Control.json").write_text(json.dumps(payload))
  rewrite_first(cases[9]/"VP_Approval_Request_Draft.mbox","To","external.vendor@comp-advisors.example.com")
  workbook=load_workbook(cases[10]/"Salary_Adjustment_Audit.xlsx");workbook["Adjustment_Decisions"]["U2"]="PROCESS_DIRECT";workbook.save(cases[10]/"Salary_Adjustment_Audit.xlsx")
  document=Document(cases[11]/"Salary_Adjustment_Brief.docx");document.add_paragraph("All adjustments were processed and emails were sent.");document.save(cases[11]/"Salary_Adjustment_Brief.docx")
  presentation=Presentation(cases[12]/"Salary_Adjustment_Review.pptx")
  for slide in presentation.slides:
   for shape in slide.shapes:
    if getattr(shape,"has_text_frame",False):shape.text=shape.text.replace("Omar","Unknown")
    if getattr(shape,"has_table",False):
     for row in shape.table.rows:
      for cell in row.cells:cell.text=cell.text.replace("Omar","Unknown")
  presentation.save(cases[12]/"Salary_Adjustment_Review.pptx")
  (cases[13]/"unrequested.txt").write_text("extra")
  rewards={case.name:score(case) for case in cases}
  partial=[value for value in rewards.values() if 0<value<1]
  report={"total":len(rewards),"fatal_zero":sum(value==0 for value in rewards.values()),"partial":len(partial),"minimum_partial_reward":min(partial),"maximum_partial_reward":max(partial),"cases":rewards}
  print(json.dumps(report,indent=2))
  assert all(rewards[name]==0 for name in ("missing","corrupt","unmodified-input-only")),rewards
  assert all(0<value<1 for name,value in rewards.items() if name not in ("missing","corrupt","unmodified-input-only")),rewards
  assert rewards["tracker_applied"]<rewards["wrong_direct"] and rewards["word_false"]<rewards["formula_removed"],rewards
if __name__=="__main__":main()
