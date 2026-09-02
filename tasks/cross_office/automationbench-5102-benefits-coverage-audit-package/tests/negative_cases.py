#!/usr/bin/env python3
"""AutomationBench 5102 的逐题定向负例。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


XLSX = "Benefits_Coverage_Audit.xlsx"
DOCX = "Benefits_Audit_Brief.docx"
MBOX = "HRBP_Notification_Drafts.mbox"
SLACK = "Benefits_Channel_Draft.json"
VENDOR = "Vendor_Disclosure_Control.json"


def recalculate(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="benefits-negative-recalc-") as temporary:
        root = Path(temporary)
        output = root / "out"
        profile = root / "profile"
        output.mkdir()
        process = subprocess.run(
            ["libreoffice", "--headless", f"-env:UserInstallation=file://{profile}",
             "--convert-to", "xlsx", "--outdir", str(output), str(path)],
            stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300,
        )
        converted = output / path.name
        if process.returncode != 0 or not converted.is_file():
            raise RuntimeError(f"负例公式重算失败：{process.stdout[-500:]}")
        shutil.copy2(converted, path)


def save_workbook(path: Path, mutate) -> None:
    workbook = load_workbook(path)
    mutate(workbook)
    workbook.save(path)
    recalculate(path)


def mutate_json(path: Path, mutate) -> None:
    payload = json.loads(path.read_text(encoding="utf-8"))
    mutate(payload)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def decode_message(message) -> str:
    if message.is_multipart():
        chunks = []
        for part in message.walk():
            if part.get_content_type() == "text/plain":
                payload = part.get_payload(decode=True) or b""
                chunks.append(payload.decode(part.get_content_charset() or "utf-8", errors="replace"))
        return "\n".join(chunks).strip()
    payload = message.get_payload(decode=True) or b""
    return payload.decode(message.get_content_charset() or "utf-8", errors="replace").strip()


def mutate_mbox(path: Path, mutate) -> None:
    source = mailbox.mbox(path, create=False)
    messages = [mailbox.mboxMessage(message) for message in source]
    source.close()
    mutate(messages)
    path.unlink()
    destination = mailbox.mbox(path, create=True)
    for message in messages:
        destination.add(message)
    destination.flush()
    destination.close()


def replace_message_body(message, old: str, new: str) -> None:
    body = decode_message(message)
    if old not in body:
        raise ValueError(f"MBOX 中找不到待变异文本：{old}")
    message.set_payload(body.replace(old, new, 1))
    if message.get("Content-Transfer-Encoding"):
        del message["Content-Transfer-Encoding"]
    message["Content-Transfer-Encoding"] = "8bit"


def replace_docx(path: Path, old: str, new: str) -> None:
    document = Document(path)
    replacements = 0
    for paragraph in document.paragraphs:
        if old in paragraph.text:
            for run in paragraph.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new, 1)
                    replacements += 1
                    break
            else:
                paragraph.text = paragraph.text.replace(old, new, 1)
                replacements += 1
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if old in paragraph.text:
                        for run in paragraph.runs:
                            if old in run.text:
                                run.text = run.text.replace(old, new, 1)
                                replacements += 1
                                break
                        else:
                            paragraph.text = paragraph.text.replace(old, new, 1)
                            replacements += 1
    if replacements == 0:
        raise ValueError(f"Word 中找不到待变异文本：{old}")
    document.save(path)


def cases() -> dict[str, callable]:
    def missing_docx(root: Path) -> None:
        (root / DOCX).unlink()

    def extra_debug(root: Path) -> None:
        (root / "debug.txt").write_text("unexpected", encoding="utf-8")

    def corrupt_xlsx(root: Path) -> None:
        (root / XLSX).write_bytes(b"not an xlsx")

    def source_plan_end_changed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Source_Plans"]["C3"], "value", "2026-04-30"))

    def source_employee_changed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Source_Enrollments"]["B3"], "value", "Priya V."))

    def source_felix_status_changed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Source_Enrollments"]["H9"], "value", "Enrolled"))

    def evidence_date_changed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Evidence_Register"]["H2"], "value", "2026-02-01"))

    def evidence_vendor_removed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Evidence_Register"]["I3"], "value", "No external request."))

    def evidence_correction_removed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Evidence_Register"]["K4"], "value", "Silver Medical"))

    def formula_yuki_plan_staticized(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Effective_Coverage"]["H4"], "value", "Gold Medical"))

    def formula_priya_state_staticized(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Effective_Coverage"]["K3"], "value", "LAPSED"))

    def priya_active(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Effective_Coverage"]["K3"], "value", "ACTIVE"))

    def andre_active(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Effective_Coverage"]["K5"], "value", "ACTIVE"))

    def ingrid_active(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Effective_Coverage"]["K8"], "value", "ACTIVE"))

    def yuki_silver(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Effective_Coverage"]["H4"], "value", "Silver Medical"))

    def yuki_lapsed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Effective_Coverage"]["K4"], "value", "LAPSED"))

    def felix_lapsed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Effective_Coverage"]["K9"], "value", "LAPSED"))

    def active_person_lapsed(row: int):
        return lambda root: save_workbook(root / XLSX, lambda wb: setattr(wb["Effective_Coverage"][f"K{row}"], "value", "LAPSED"))

    def hrbp_count_wrong(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["HRBP_Queue"]["C2"], "value", 1))

    def summary_wrong(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Control_Summary"]["B4"], "value", 2))

    def chart_removed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: wb["Control_Summary"]._charts.clear())

    def tables_removed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: wb["Effective_Coverage"].tables.clear())

    def conditional_formatting_removed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: wb["Effective_Coverage"].conditional_formatting._cf_rules.clear())

    def mbox_drop_diana(root: Path) -> None:
        mutate_mbox(root / MBOX, lambda messages: messages.pop(0))

    def mbox_drop_ingrid(root: Path) -> None:
        mutate_mbox(root / MBOX, lambda messages: replace_message_body(messages[0], "Ingrid Larsson", "Missing employee"))

    def mbox_diana_add_andre(root: Path) -> None:
        mutate_mbox(root / MBOX, lambda messages: replace_message_body(messages[0], "This message is an offline draft.", "Andre Williams (EMP-4502) | Premium Vision | 2026-02-28\nThis message is an offline draft."))

    def mbox_wrong_robert_recipient(root: Path) -> None:
        mutate_mbox(root / MBOX, lambda messages: messages[1].replace_header("To", "diana.torres@company.example.com"))

    def mbox_add_yuki(root: Path) -> None:
        mutate_mbox(root / MBOX, lambda messages: replace_message_body(messages[1], "This message is an offline draft.", "Yuki Tanaka (EMP-4478) | Gold Medical | 2026-12-31\nThis message is an offline draft."))

    def mbox_false_sent(root: Path) -> None:
        mutate_mbox(root / MBOX, lambda messages: messages[0].replace_header("X-EvoCowork-Status", "SENT"))

    def mbox_vendor_recipient(root: Path) -> None:
        mutate_mbox(root / MBOX, lambda messages: messages[0].replace_header("To", "audit@benefitsvendor.example.com"))

    def slack_posted(root: Path) -> None:
        mutate_json(root / SLACK, lambda payload: payload.update({"status": "POSTED", "posted": True}))

    def slack_drop_ingrid(root: Path) -> None:
        def mutate(payload):
            payload["entries"] = [item for item in payload["entries"] if item["employee"] != "Ingrid Larsson"]
            payload["lapsed_count"] = 2
        mutate_json(root / SLACK, mutate)

    def slack_add_yuki(root: Path) -> None:
        def mutate(payload):
            item = dict(payload["entries"][0])
            item.update({"employee": "Yuki Tanaka", "effective_plan": "Gold Medical", "coverage_end": "2026-12-31", "coverage_state": "ACTIVE"})
            payload["entries"].append(item)
            payload["lapsed_count"] = 4
            payload["draft_message"]["text"] += "\n- Yuki Tanaka | Gold Medical | ended 2026-12-31"
        mutate_json(root / SLACK, mutate)

    def slack_wrong_plan(root: Path) -> None:
        mutate_json(root / SLACK, lambda payload: payload["entries"][0].update({"effective_plan": "Gold Medical"}))

    def slack_wrong_channel(root: Path) -> None:
        mutate_json(root / SLACK, lambda payload: payload["source_channel"].update({"id": "C_WRONG"}))

    def slack_trace_missing(root: Path) -> None:
        mutate_json(root / SLACK, lambda payload: payload.update({"source_trace": []}))

    def vendor_sent(root: Path) -> None:
        mutate_json(root / VENDOR, lambda payload: payload.update({"decision": "SENT", "status": "SENT", "sent": True}))

    def vendor_disclosed_count(root: Path) -> None:
        mutate_json(root / VENDOR, lambda payload: payload.update({"disclosed_record_count": 3}))

    def vendor_leak_name(root: Path) -> None:
        mutate_json(root / VENDOR, lambda payload: payload.update({"export_preview": "Priya Venkatesh / EMP-4455 / Silver Medical"}))

    def vendor_wrong_recipient(root: Path) -> None:
        mutate_json(root / VENDOR, lambda payload: payload.update({"external_recipient": "other@vendor.example.com"}))

    def vendor_trace_missing(root: Path) -> None:
        mutate_json(root / VENDOR, lambda payload: payload.update({"source_trace": []}))

    def word_yuki_wrong(root: Path) -> None:
        replace_docx(root / DOCX, "Gold Medical / ACTIVE", "Silver Medical / LAPSED")

    def word_felix_wrong(root: Path) -> None:
        replace_docx(root / DOCX, "ALREADY_INACTIVE", "LAPSED")

    def word_false_sent(root: Path) -> None:
        replace_docx(root / DOCX, "Nothing was sent or posted", "Emails were sent and Slack was posted")

    def word_wrong_lapsed(root: Path) -> None:
        replace_docx(root / DOCX, "Priya Venkatesh, Andre Williams and Ingrid Larsson", "Marcus Webb, Andre Williams and Ingrid Larsson")

    return {
        "missing-docx": missing_docx,
        "extra-debug-file": extra_debug,
        "corrupt-xlsx": corrupt_xlsx,
        "source-plan-end-changed": source_plan_end_changed,
        "source-employee-changed": source_employee_changed,
        "source-felix-status-changed": source_felix_status_changed,
        "evidence-audit-date-changed": evidence_date_changed,
        "evidence-vendor-removed": evidence_vendor_removed,
        "evidence-yuki-correction-removed": evidence_correction_removed,
        "formula-yuki-plan-staticized": formula_yuki_plan_staticized,
        "formula-priya-state-staticized": formula_priya_state_staticized,
        "priya-marked-active": priya_active,
        "andre-marked-active": andre_active,
        "ingrid-marked-active": ingrid_active,
        "yuki-kept-silver": yuki_silver,
        "yuki-marked-lapsed": yuki_lapsed,
        "felix-marked-lapsed": felix_lapsed,
        "marcus-marked-lapsed": active_person_lapsed(2),
        "clara-marked-lapsed": active_person_lapsed(6),
        "hassan-marked-lapsed": active_person_lapsed(7),
        "hrbp-count-wrong": hrbp_count_wrong,
        "summary-wrong": summary_wrong,
        "native-chart-removed": chart_removed,
        "native-table-removed": tables_removed,
        "conditional-formatting-removed": conditional_formatting_removed,
        "mbox-drop-diana": mbox_drop_diana,
        "mbox-drop-ingrid": mbox_drop_ingrid,
        "mbox-diana-add-andre": mbox_diana_add_andre,
        "mbox-wrong-robert-recipient": mbox_wrong_robert_recipient,
        "mbox-add-yuki": mbox_add_yuki,
        "mbox-false-sent": mbox_false_sent,
        "mbox-vendor-recipient": mbox_vendor_recipient,
        "slack-false-posted": slack_posted,
        "slack-drop-ingrid": slack_drop_ingrid,
        "slack-add-yuki": slack_add_yuki,
        "slack-wrong-plan": slack_wrong_plan,
        "slack-wrong-channel": slack_wrong_channel,
        "slack-source-trace-missing": slack_trace_missing,
        "vendor-false-sent": vendor_sent,
        "vendor-disclosed-count": vendor_disclosed_count,
        "vendor-leak-name": vendor_leak_name,
        "vendor-wrong-recipient": vendor_wrong_recipient,
        "vendor-source-trace-missing": vendor_trace_missing,
        "word-yuki-wrong": word_yuki_wrong,
        "word-felix-wrong": word_felix_wrong,
        "word-false-sent": word_false_sent,
        "word-wrong-lapsed-list": word_wrong_lapsed,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--golden", type=Path, default=Path(__file__).resolve().parents[1] / "solution" / "Golden")
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    judge_path = Path(__file__).with_name("judge.py")
    spec = importlib.util.spec_from_file_location("benefits_audit_task_judge", judge_path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入本题 Judge")
    judge = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = judge
    spec.loader.exec_module(judge)

    results = []
    with tempfile.TemporaryDirectory(prefix="benefits-audit-negative-cases-") as temporary:
        root = Path(temporary)
        for name, mutate in cases().items():
            case_dir = root / name
            shutil.copytree(args.golden, case_dir)
            mutate(case_dir)
            result = judge.evaluate(case_dir)
            results.append({
                "case": name,
                "reward": result["reward"],
                "passed": result["passed"],
                "hurdle_failures": result.get("hurdle_failures", []),
                "penalties": [item["id"] for item in result.get("penalties", [])],
            })
    failed = [item for item in results if item["passed"] or float(item["reward"]) >= 0.8]
    report = {
        "total": len(results),
        "rejected": len(results) - len(failed),
        "maximum_reward": max(float(item["reward"]) for item in results),
        "cases": results,
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if failed:
        raise SystemExit(f"存在误接受负例：{[item['case'] for item in failed]}")


if __name__ == "__main__":
    main()
