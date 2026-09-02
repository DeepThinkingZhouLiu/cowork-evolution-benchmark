"""Office 交付物的确定性渲染基础设施；不包含任何任务业务真值。"""

from __future__ import annotations

import re
import zipfile
from datetime import datetime, timezone
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt as PptPt


FIXED_TIME = datetime(2026, 4, 15, 9, 0, tzinfo=timezone.utc)
ZIP_TIME = (2026, 4, 15, 9, 0, 0)
NAVY, BLUE, PALE, GREEN, AMBER, RED, GRAY, WHITE = "17365D", "4472C4", "D9EAF7", "E2F0D9", "FFF2CC", "F4CCCC", "E7E6E6", "FFFFFF"
THIN = Side(style="thin", color="B7C9D6")


def normalize_package(path: Path) -> None:
    temporary = path.with_suffix(path.suffix + ".normalized")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            payload = source.read(name)
            if name == "docProps/core.xml":
                payload = re.sub(rb"(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)", rb"\g<1>2026-04-15T09:00:00Z\g<2>", payload)
            info = zipfile.ZipInfo(name, ZIP_TIME); info.compress_type = zipfile.ZIP_DEFLATED; info.external_attr = 0o600 << 16
            target.writestr(info, payload)
    temporary.replace(path)


def _style_sheet(ws, widths: list[int] | None = None) -> None:
    ws.sheet_view.showGridLines = False; ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY); cell.font = Font(name="Carlito", bold=True, color=WHITE, size=9)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = Border(bottom=THIN)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Carlito", size=8); cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color="D9E2F3"))
    if widths:
        for index, width in enumerate(widths, 1): ws.column_dimensions[ws.cell(1, index).column_letter].width = width
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.print_area = ws.dimensions


def write_workbook(path: Path, title: str, sheets: list[dict[str, Any]]) -> None:
    wb = Workbook(); wb.remove(wb.active); wb.properties.title = title; wb.properties.creator = "EvoCowork reference solution"
    wb.properties.created = FIXED_TIME.replace(tzinfo=None); wb.properties.modified = FIXED_TIME.replace(tzinfo=None)
    for spec in sheets:
        ws = wb.create_sheet(spec["name"]); ws.append(spec["headers"])
        for row in spec.get("rows", []): ws.append(row)
        _style_sheet(ws, spec.get("widths"))
        if ws.max_row >= 2:
            table = Table(displayName=spec["table"], ref=ws.dimensions); table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True); ws.add_table(table)
        for rule in spec.get("conditional", []):
            ws.conditional_formatting.add(rule["range"], FormulaRule(formula=[rule["formula"]], fill=PatternFill("solid", fgColor=rule["color"])))
        chart_spec = spec.get("chart")
        if chart_spec:
            chart = BarChart(); chart.title = chart_spec["title"]; chart.y_axis.title = chart_spec.get("axis", "Count")
            chart.add_data(Reference(ws, min_col=chart_spec["value_col"], min_row=chart_spec["min_row"], max_row=chart_spec["max_row"]), titles_from_data=False)
            chart.set_categories(Reference(ws, min_col=chart_spec["label_col"], min_row=chart_spec["min_row"], max_row=chart_spec["max_row"])); chart.height = 6; chart.width = 10; ws.add_chart(chart, chart_spec.get("anchor", "E2"))
    wb.calculation.fullCalcOnLoad = True; wb.calculation.forceFullCalc = True; wb.calculation.calcMode = "auto"
    wb.save(path); normalize_package(path)


def _shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr(); shading = OxmlElement("w:shd"); shading.set(qn("w:fill"), color); tc_pr.append(shading)


def _doc_table(doc: Document, headers: list[str], rows: list[list[Any]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers)); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.style = "Table Grid"
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]; cell.text = str(value); _shade(cell, NAVY)
        for run in cell.paragraphs[0].runs: run.font.bold = True; run.font.color.rgb = RGBColor(255, 255, 255); run.font.size = Pt(8)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value); cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for paragraph in cells[index].paragraphs:
                for run in paragraph.runs: run.font.name = "Carlito"; run.font.size = Pt(8)


def write_docx(path: Path, title: str, subtitle: str, pages: list[dict[str, Any]]) -> None:
    doc = Document(); section = doc.sections[0]; section.top_margin = Cm(1.4); section.bottom_margin = Cm(1.4); section.left_margin = Cm(1.5); section.right_margin = Cm(1.5)
    doc.styles["Normal"].font.name = "Carlito"; doc.styles["Normal"].font.size = Pt(9)
    heading = doc.add_heading(title, 0); heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph = doc.add_paragraph(subtitle); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for page_index, page in enumerate(pages):
        if page_index: doc.add_section(WD_SECTION.NEW_PAGE)
        for block in page.get("blocks", []):
            kind = block["type"]
            if kind == "heading": doc.add_heading(block["text"], level=block.get("level", 1))
            elif kind == "paragraph": doc.add_paragraph(block["text"])
            elif kind == "table": _doc_table(doc, block["headers"], block["rows"])
    doc.core_properties.author = "EvoCowork reference solution"; doc.core_properties.created = FIXED_TIME; doc.core_properties.modified = FIXED_TIME
    doc.save(path); normalize_package(path)


def _ppt_table(slide, headers: list[str], rows: list[list[Any]], top: float = 1.48) -> None:
    table = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(0.45), Inches(top), Inches(12.35), Inches(5.45)).table
    for col, header in enumerate(headers):
        cell = table.cell(0, col); cell.text = str(header); cell.fill.solid(); cell.fill.fore_color.rgb = PptRGB(23, 54, 93)
    for row_index, values in enumerate(rows, 1):
        for col, value in enumerate(values):
            cell = table.cell(row_index, col); cell.text = str(value); cell.fill.solid(); cell.fill.fore_color.rgb = PptRGB(242, 246, 250) if row_index % 2 == 0 else PptRGB(255, 255, 255)
    for row in range(len(rows) + 1):
        for col in range(len(headers)):
            for paragraph in table.cell(row, col).text_frame.paragraphs:
                paragraph.font.name = "Carlito"; paragraph.font.size = PptPt(8 if len(rows) > 6 else 10); paragraph.alignment = PP_ALIGN.LEFT
                if row == 0: paragraph.font.bold = True; paragraph.font.color.rgb = PptRGB(255, 255, 255)


def write_pptx(path: Path, slides: list[dict[str, Any]]) -> None:
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5); blank = prs.slide_layouts[6]
    for spec in slides:
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(0.45), Inches(0.28), Inches(12.4), Inches(0.78)); paragraph = box.text_frame.paragraphs[0]
        paragraph.text = spec["title"]; paragraph.font.name = "Carlito"; paragraph.font.size = PptPt(25); paragraph.font.bold = True; paragraph.font.color.rgb = PptRGB(23, 54, 93)
        sub = slide.shapes.add_textbox(Inches(0.48), Inches(1.0), Inches(12.2), Inches(0.38)); paragraph = sub.text_frame.paragraphs[0]
        paragraph.text = spec["subtitle"]; paragraph.font.name = "Carlito"; paragraph.font.size = PptPt(10); paragraph.font.color.rgb = PptRGB(89, 89, 89)
        _ppt_table(slide, spec["headers"], spec["rows"])
    prs.core_properties.author = "EvoCowork reference solution"; prs.core_properties.created = FIXED_TIME; prs.core_properties.modified = FIXED_TIME
    prs.save(path); normalize_package(path)


def message_bytes(draft: dict[str, Any], order: int) -> bytes:
    message = EmailMessage(policy=SMTP); message["From"] = draft.get("from", "hr-ops@company.example.com"); message["To"] = draft["to"]
    if draft.get("cc"): message["Cc"] = draft["cc"]
    message["Subject"] = draft["subject"]; message["Date"] = "Wed, 15 Apr 2026 09:00:00 +0000"
    message["Message-ID"] = f"<{draft.get('id_prefix', 'draft')}-{order:02d}@draft.invalid>"; message["X-EvoCowork-State"] = draft.get("state", "DRAFT_NOT_SENT"); message["X-EvoCowork-Type"] = draft["type"]
    message.set_content(draft["body"], charset="utf-8"); buffer = BytesIO(); BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message); return buffer.getvalue()


def write_mbox(path: Path, drafts: list[dict[str, Any]]) -> None:
    chunks: list[bytes] = []
    for order, draft in enumerate(drafts, 1):
        chunks.append(f"From {draft.get('from', 'hr-ops@company.example.com')} Wed Apr 15 09:00:{order:02d} 2026\n".encode("ascii")); chunks.append(message_bytes(draft, order).replace(b"\r\n", b"\n")); chunks.append(b"\n")
    path.write_bytes(b"".join(chunks))
