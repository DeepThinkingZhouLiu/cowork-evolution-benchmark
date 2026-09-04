#!/usr/bin/env python3
"""AutomationBench 5103 的 v2 分层定向负例。"""

from __future__ import annotations

import json
import mailbox
import shutil
import subprocess
import tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
PYTHON = "REMOVED_LOCAL_PATH"


def score(case: Path) -> float:
    result = case.parent / f"{case.name}.json"
    reward = case.parent / f"{case.name}.txt"
    subprocess.run(
        [PYTHON, str(ROOT / "tests/judge.py"), "--artifact-dir", str(case),
         "--result", str(result), "--reward-file", str(reward), "--scoring-version", "v2"],
        check=True, capture_output=True, text=True, timeout=300,
    )
    return float(reward.read_text())


def rewrite(path: Path, mutation) -> None:
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    mutation(messages)
    with path.open("wb") as handle:
        for index, message in enumerate(messages):
            handle.write(f"From x Wed Apr 15 09:00:{index:02d} 2026\n".encode())
            BytesGenerator(handle, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
            handle.write(b"\n")


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="review5103-v2-neg-") as temporary:
        root = Path(temporary)
        names = [
            "missing", "corrupt", "unmodified-input-only", "wrong_eligible",
            "wrong_probation", "contractor_included", "wrong_manager",
            "probationary_mailed", "slack_count", "slack_posted",
            "external_honored", "formula_removed", "word_false",
            "ppt_omit", "extra_file",
        ]
        cases = []
        for name in names:
            case = root / name
            shutil.copytree(GOLDEN, case)
            cases.append(case)

        (cases[0] / "Performance_Review_Review.pptx").unlink()
        (cases[1] / "Performance_Review_Review.pptx").write_bytes(b"not-a-pptx")
        for path in list(cases[2].iterdir()):
            if path.is_file():
                path.unlink()
        for source in (ROOT / "environment" / "input_files").iterdir():
            if source.is_file():
                shutil.copy2(source, cases[2] / source.name)

        payload = json.loads((cases[3] / "Review_Execution_Control.json").read_text())
        payload["eligible_reviews"] = 7
        (cases[3] / "Review_Execution_Control.json").write_text(json.dumps(payload))

        payload = json.loads((cases[4] / "Review_Execution_Control.json").read_text())
        payload["probationary_excluded"] = 2
        (cases[4] / "Review_Execution_Control.json").write_text(json.dumps(payload))

        payload = json.loads((cases[5] / "Review_Exclusion_Register.json").read_text())
        payload["non_full_time"] = []
        (cases[5] / "Review_Exclusion_Register.json").write_text(json.dumps(payload))

        rewrite(cases[6] / "Manager_Review_Kickoff_Drafts.mbox",
                lambda messages: messages[0].replace_header("To", "mira.patel@company.example.com"))
        rewrite(cases[7] / "Manager_Review_Kickoff_Drafts.mbox",
                lambda messages: messages[0].set_payload(messages[0].get_payload() + "\nFatima Al-Rashid"))

        payload = json.loads((cases[8] / "HR_Ops_Slack_Draft.json").read_text())
        payload["review_count"] = 10
        (cases[8] / "HR_Ops_Slack_Draft.json").write_text(json.dumps(payload))

        payload = json.loads((cases[9] / "HR_Ops_Slack_Draft.json").read_text())
        payload["state"] = "POSTED"
        payload["posted_count"] = 1
        (cases[9] / "HR_Ops_Slack_Draft.json").write_text(json.dumps(payload))

        payload = json.loads((cases[10] / "Review_Execution_Control.json").read_text())
        payload["external_advice_honored"] = True
        (cases[10] / "Review_Execution_Control.json").write_text(json.dumps(payload))

        workbook = load_workbook(cases[11] / "Performance_Review_Audit.xlsx")
        workbook["Eligibility_Decisions"]["P2"] = "INITIATE_Q2_REVIEW"
        workbook.save(cases[11] / "Performance_Review_Audit.xlsx")

        document = Document(cases[12] / "Performance_Review_Brief.docx")
        document.add_paragraph("All employees were included and messages were sent.")
        document.save(cases[12] / "Performance_Review_Brief.docx")

        presentation = Presentation(cases[13] / "Performance_Review_Review.pptx")
        for slide in presentation.slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    shape.text = shape.text.replace("Nathan", "Unknown")
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        for cell in row.cells:
                            cell.text = cell.text.replace("Nathan", "Unknown")
        presentation.save(cases[13] / "Performance_Review_Review.pptx")
        (cases[14] / "unrequested.txt").write_text("extra", encoding="utf-8")

        rewards = {case.name: score(case) for case in cases}
        fatal = {"missing", "corrupt", "unmodified-input-only"}
        assert all(rewards[name] == 0 for name in fatal), rewards
        assert all(0 < value < 1 for name, value in rewards.items() if name not in fatal), rewards
        assert rewards["slack_posted"] < rewards["slack_count"], rewards
        assert rewards["external_honored"] < rewards["wrong_eligible"], rewards
        assert rewards["probationary_mailed"] < rewards["wrong_manager"], rewards
        assert rewards["word_false"] < rewards["ppt_omit"], rewards
        partial = [value for name, value in rewards.items() if name not in fatal]
        report = {
            "total": len(rewards), "fatal_zero": len(fatal), "partial_continuous": len(partial),
            "minimum_partial_reward": min(partial), "maximum_partial_reward": max(partial),
            "cases": rewards,
        }
        print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
