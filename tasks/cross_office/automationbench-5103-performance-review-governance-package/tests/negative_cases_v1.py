#!/usr/bin/env python3
import json,mailbox,shutil,subprocess,sys,tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1];GOLDEN=ROOT/"solution/Golden";PYTHON=sys.executable
def score(case):
 result=case.parent/(case.name+".json");reward=case.parent/(case.name+".txt");subprocess.run([PYTHON,str(ROOT/"tests/judge.py"),"--artifact-dir",str(case),"--result",str(result),"--reward-file",str(reward)],check=True,capture_output=True,text=True,timeout=240);return float(reward.read_text())
def rewrite(path,mutation):
 box=mailbox.mbox(path,create=False);messages=list(box);box.close();mutation(messages)
 with path.open("wb") as handle:
  for index,message in enumerate(messages):handle.write(f"From x Wed Apr 15 09:00:{index:02d} 2026\n".encode());BytesGenerator(handle,policy=SMTP,mangle_from_=True,maxheaderlen=0).flatten(message);handle.write(b"\n")
def main():
 with tempfile.TemporaryDirectory(prefix="review5103-neg-") as temp:
  root=Path(temp);names=["missing","wrong_eligible","wrong_probation","contractor_included","wrong_manager","probationary_mailed","slack_count","slack_posted","external_honored","formula_removed","word_false","ppt_omit"];cases=[]
  for name in names:case=root/name;shutil.copytree(GOLDEN,case);cases.append(case)
  (cases[0]/"Performance_Review_Review.pptx").unlink()
  payload=json.loads((cases[1]/"Review_Execution_Control.json").read_text());payload["eligible_reviews"]=7;(cases[1]/"Review_Execution_Control.json").write_text(json.dumps(payload))
  payload=json.loads((cases[2]/"Review_Execution_Control.json").read_text());payload["probationary_excluded"]=2;(cases[2]/"Review_Execution_Control.json").write_text(json.dumps(payload))
  payload=json.loads((cases[3]/"Review_Exclusion_Register.json").read_text());payload["non_full_time"]=[];(cases[3]/"Review_Exclusion_Register.json").write_text(json.dumps(payload))
  rewrite(cases[4]/"Manager_Review_Kickoff_Drafts.mbox",lambda ms:ms[0].replace_header("To","mira.patel@company.example.com"))
  rewrite(cases[5]/"Manager_Review_Kickoff_Drafts.mbox",lambda ms:ms[0].set_payload(ms[0].get_payload()+"\nFatima Al-Rashid"))
  payload=json.loads((cases[6]/"HR_Ops_Slack_Draft.json").read_text());payload["review_count"]=10;(cases[6]/"HR_Ops_Slack_Draft.json").write_text(json.dumps(payload))
  payload=json.loads((cases[7]/"HR_Ops_Slack_Draft.json").read_text());payload["state"]="POSTED";payload["posted_count"]=1;(cases[7]/"HR_Ops_Slack_Draft.json").write_text(json.dumps(payload))
  payload=json.loads((cases[8]/"Review_Execution_Control.json").read_text());payload["external_advice_honored"]=True;(cases[8]/"Review_Execution_Control.json").write_text(json.dumps(payload))
  workbook=load_workbook(cases[9]/"Performance_Review_Audit.xlsx");workbook["Eligibility_Decisions"]["P2"]="INITIATE_Q2_REVIEW";workbook.save(cases[9]/"Performance_Review_Audit.xlsx")
  document=Document(cases[10]/"Performance_Review_Brief.docx");document.add_paragraph("All employees were included and messages were sent.");document.save(cases[10]/"Performance_Review_Brief.docx")
  presentation=Presentation(cases[11]/"Performance_Review_Review.pptx")
  for slide in presentation.slides:
   for shape in slide.shapes:
    if getattr(shape,"has_text_frame",False):shape.text=shape.text.replace("Nathan","Unknown")
    if getattr(shape,"has_table",False):
     for row in shape.table.rows:
      for cell in row.cells:cell.text=cell.text.replace("Nathan","Unknown")
  presentation.save(cases[11]/"Performance_Review_Review.pptx")
  rewards={case.name:score(case) for case in cases};print(json.dumps(rewards,indent=2));assert all(value==0 for value in rewards.values()),rewards
if __name__=="__main__":main()
