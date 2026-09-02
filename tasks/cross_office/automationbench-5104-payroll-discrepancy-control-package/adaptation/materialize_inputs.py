#!/usr/bin/env python3
"""只把 AutomationBench 5104 的公开初始状态静态化为本题输入。

本脚本只接受 ``hr.payroll_discrepancy``，并逐项断言八个工资差异、
四条升级政策、一封纠错邮件、频道、时间和十五条上游断言。它不是
通用制题器，也不会创造新的员工、金额或业务规则。
"""

from __future__ import annotations

import importlib.util
import json
import shutil
import sys
import types
import zipfile
from datetime import datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
UPSTREAM = Path(__file__).with_name("upstream-hr-tasks.py")
SNAPSHOT = Path(__file__).with_name("upstream-task-5104.json")
DATA_DIR = ROOT / "data" / "input_files"
ENV_DIR = ROOT / "environment" / "input_files"
FIXED_ZIP_TIME = (2026, 4, 15, 9, 0, 0)


def load_task() -> dict:
    datasets = types.ModuleType("datasets")
    datasets.Dataset = object
    sys.modules["datasets"] = datasets
    for name in ("automationbench", "automationbench.domains", "automationbench.domains.hr"):
        sys.modules[name] = types.ModuleType(name)
    noise = types.ModuleType("automationbench.domains.hr._noise")
    noise.apply_noise = lambda value: value
    sys.modules[noise.__name__] = noise

    spec = importlib.util.spec_from_file_location("automationbench_hr_tasks_frozen_5104", UPSTREAM)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入冻结的 AutomationBench HR 源文件")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    task = module.get_task_5104()
    if task.get("example_id") != 5104 or task.get("task") != "hr.payroll_discrepancy":
        raise ValueError("冻结函数不再是 AutomationBench 5104")
    return task


def assert_shape(task: dict) -> None:
    state = task["info"]["initial_state"]
    assert state["meta"] == {"schema_version": "0.1.0", "current_time": "2026-04-15T09:00:00Z"}
    assert [item["id"] for item in state["gmail"]["messages"]] == ["msg_payroll_correction_5104"]
    source = state["google_sheets"]["spreadsheets"]
    assert len(source) == 1 and source[0]["id"] == "ss_payroll_5104"
    worksheets = source[0]["worksheets"]
    assert [item["title"] for item in worksheets] == ["April Discrepancies", "Escalation Policy"]
    assert [item["row_id"] for item in worksheets[0]["rows"]] == list(range(2, 10))
    assert [item["cells"]["Employee"] for item in worksheets[0]["rows"]] == [
        "Elise Fontaine", "George Nakamura", "Hannah Olsen", "Ibrahim Sayed",
        "Jenny Park", "Kevin Obi", "Lina Marchetti", "Miguel Santos",
    ]
    assert [item["row_id"] for item in worksheets[1]["rows"]] == [2, 3, 4, 5]
    assert [item["cells"]["Rule"] for item in worksheets[1]["rows"]] == [
        "Variance <= $500", "Variance > $500", "Disputed/On Hold", "Direct Back-Pay",
    ]
    assert state["slack"]["channels"] == [
        {"id": "C_PAY_5104", "name": "payroll", "is_private": False}
    ]
    assert state["slack"]["messages"] == []
    assert len(task["info"]["assertions"]) == 15


def normalize_xlsx(path: Path) -> None:
    temporary = path.with_suffix(".normalized.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        temporary, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as destination:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            destination.writestr(info, source.read(name))
    temporary.replace(path)


def style_source_sheet(sheet, widths: dict[str, float]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = sheet.dimensions


def write_tracker(state: dict, target: Path) -> None:
    spreadsheet = state["google_sheets"]["spreadsheets"][0]
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.creator = "AutomationBench 5104 public fixture"
    workbook.title = spreadsheet["title"]
    workbook.subject = "Static Office representation of payroll discrepancy and escalation sources"
    workbook.description = "No employee, amount, note or policy was added beyond AutomationBench 5104."

    discrepancy_sheet = workbook.create_sheet("April Discrepancies")
    discrepancy_headers = [
        "Source Row", "Employee", "Employee Email", "Manager Email", "Issue",
        "Recorded Amount", "Notes", "Status", "Variance",
    ]
    discrepancy_sheet.append(discrepancy_headers)
    for item in spreadsheet["worksheets"][0]["rows"]:
        cells = item["cells"]
        discrepancy_sheet.append(
            [item["row_id"]] + [cells.get(header, "") for header in discrepancy_headers[1:]]
        )
    discrepancy_table = Table(displayName="PayrollDiscrepancySourceInput", ref=discrepancy_sheet.dimensions)
    discrepancy_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    discrepancy_sheet.add_table(discrepancy_table)
    style_source_sheet(discrepancy_sheet, {
        "A": 12, "B": 22, "C": 34, "D": 34, "E": 30,
        "F": 18, "G": 52, "H": 14, "I": 14,
    })

    policy_sheet = workbook.create_sheet("Escalation Policy")
    policy_headers = ["Source Row", "Rule", "Action"]
    policy_sheet.append(policy_headers)
    for item in spreadsheet["worksheets"][1]["rows"]:
        cells = item["cells"]
        policy_sheet.append([item["row_id"], cells["Rule"], cells["Action"]])
    policy_table = Table(displayName="PayrollEscalationPolicyInput", ref=policy_sheet.dimensions)
    policy_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    policy_sheet.add_table(policy_table)
    style_source_sheet(policy_sheet, {"A": 12, "B": 24, "C": 100})

    workbook.save(target)
    normalize_xlsx(target)


def message_bytes(record: dict) -> bytes:
    message = EmailMessage(policy=SMTP)
    message["From"] = record["from_"]
    message["To"] = ", ".join(record["to"])
    message["Subject"] = record["subject"]
    parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
    message["Date"] = parsed.strftime("%a, %d %b %Y %H:%M:%S +0000")
    message["Message-ID"] = f"<{record['id']}@automationbench.invalid>"
    message["X-AutomationBench-Message-ID"] = record["id"]
    message["X-AutomationBench-Thread-ID"] = record["thread_id"]
    message["X-AutomationBench-Labels"] = ",".join(record.get("label_ids", []))
    message.set_content(record["body_plain"], charset="utf-8")
    buffer = BytesIO()
    BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
    return buffer.getvalue()


def write_mbox(messages: list[dict], target: Path) -> None:
    chunks: list[bytes] = []
    for record in messages:
        parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
        chunks.append(f"From {record['from_']} {parsed.strftime('%a %b %d %H:%M:%S %Y')}\n".encode("ascii"))
        chunks.append(message_bytes(record).replace(b"\r\n", b"\n"))
        chunks.append(b"\n")
    target.write_bytes(b"".join(chunks))


def main() -> None:
    task = load_task()
    assert_shape(task)
    state = task["info"]["initial_state"]
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    ENV_DIR.mkdir(parents=True, exist_ok=True)

    SNAPSHOT.write_text(json.dumps(task, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_tracker(state, DATA_DIR / "Payroll_Discrepancy_Tracker.xlsx")
    write_mbox(state["gmail"]["messages"], DATA_DIR / "Payroll_Corrections_Mailbox.mbox")
    (DATA_DIR / "Payroll_Channel.json").write_text(
        json.dumps(state["slack"], ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (DATA_DIR / "Workflow_Context.json").write_text(
        json.dumps({"meta": state["meta"]}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    for source in sorted(DATA_DIR.iterdir()):
        if source.is_file():
            shutil.copy2(source, ENV_DIR / source.name)


if __name__ == "__main__":
    main()
