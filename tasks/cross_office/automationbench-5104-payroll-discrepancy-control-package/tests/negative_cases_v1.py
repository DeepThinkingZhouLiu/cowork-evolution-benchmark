#!/usr/bin/env python3
"""AutomationBench 5104 的逐题定向负例。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


XLSX = "Payroll_Discrepancy_Audit.xlsx"
DOCX = "Payroll_Control_Brief.docx"
MBOX = "Payroll_Routing_Drafts.mbox"
SLACK = "Payroll_Channel_Draft.json"
CONTROL = "Back_Pay_Execution_Control.json"


def recalculate(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="payroll-negative-recalc-") as temporary:
        root = Path(temporary)
        source_dir = root / "source"
        output_dir = root / "output"
        profile = root / "profile"
        source_dir.mkdir()
        output_dir.mkdir()
        source = source_dir / path.name
        shutil.copy2(path, source)
        process = subprocess.run(
            [
                "libreoffice", "--headless", f"-env:UserInstallation=file://{profile}",
                "--convert-to", "xlsx", "--outdir", str(output_dir), str(source),
            ],
            stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300,
        )
        converted = output_dir / path.name
        if process.returncode != 0 or not converted.is_file():
            raise RuntimeError(f"负例公式重算失败：{process.stdout[-500:]}")
        shutil.copy2(converted, path)


def save_workbook(path: Path, mutate) -> None:
    workbook = load_workbook(path)
    mutate(workbook)
    workbook.save(path)
    recalculate(path)


def mutate_json(path: Path, mutate) -> None:
    payload = json.loads(path.read_text(encoding="utf-8"))
    mutate(payload)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def decode_message(message) -> str:
    if message.is_multipart():
        chunks: list[str] = []
        for part in message.walk():
            if part.get_content_type() == "text/plain":
                payload = part.get_payload(decode=True) or b""
                chunks.append(payload.decode(part.get_content_charset() or "utf-8", errors="replace"))
        return "\n".join(chunks).strip()
    payload = message.get_payload(decode=True) or b""
    return payload.decode(message.get_content_charset() or "utf-8", errors="replace").strip()


def mutate_mbox(path: Path, mutate) -> None:
    source = mailbox.mbox(path, create=False)
    messages = [mailbox.mboxMessage(message) for message in source]
    source.close()
    mutate(messages)
    path.unlink()
    destination = mailbox.mbox(path, create=True)
    for message in messages:
        destination.add(message)
    destination.flush()
    destination.close()


def replace_message_body(message, old: str, new: str) -> None:
    body = decode_message(message)
    if old not in body:
        raise ValueError(f"MBOX 中找不到待变异文本：{old}")
    message.set_payload(body.replace(old, new, 1))
    if message.get("Content-Transfer-Encoding"):
        del message["Content-Transfer-Encoding"]
    message["Content-Transfer-Encoding"] = "8bit"


def replace_docx(path: Path, old: str, new: str) -> None:
    document = Document(path)
    replacements = 0
    for paragraph in document.paragraphs:
        if old in paragraph.text:
            for run in paragraph.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new, 1)
                    replacements += 1
                    break
            else:
                paragraph.text = paragraph.text.replace(old, new, 1)
                replacements += 1
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if old in paragraph.text:
                        for run in paragraph.runs:
                            if old in run.text:
                                run.text = run.text.replace(old, new, 1)
                                replacements += 1
                                break
                        else:
                            paragraph.text = paragraph.text.replace(old, new, 1)
                            replacements += 1
    if replacements == 0:
        raise ValueError(f"Word 中找不到待变异文本：{old}")
    document.save(path)


def cases() -> dict[str, callable]:
    def missing_docx(root: Path) -> None:
        (root / DOCX).unlink()

    def extra_debug(root: Path) -> None:
        (root / "debug.txt").write_text("unexpected", encoding="utf-8")

    def corrupt_xlsx(root: Path) -> None:
        (root / XLSX).write_bytes(b"not an xlsx")

    def source_name_changed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Source_Cases"]["B2"], "value", "Elise Fontane"))

    def source_issue_changed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Source_Cases"]["E3"], "value", "Ordinary bonus"))

    def source_recorded_changed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Source_Cases"]["F5"], "value", "$3,100"))

    def source_note_changed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Source_Cases"]["G4"], "value", "Night shift premium $150 not applied"))

    def policy_threshold_changed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Source_Policy"]["B3"], "value", "Variance > $5,000"))

    def policy_backpay_changed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Source_Policy"]["C5"], "value", "HR Ops may execute back-pay."))

    def evidence_mail_removed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Evidence_Register"]["F10"], "value", "Email omitted"))

    def evidence_elise_amount(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Evidence_Register"]["E2"], "value", 360))

    def evidence_george_amount(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Evidence_Register"]["E3"], "value", 2000))

    def evidence_channel_wrong(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Evidence_Register"]["D12"], "value", "C_WRONG"))

    def evidence_time_wrong(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Evidence_Register"]["F11"], "value", "2026-04-14T09:00:00Z"))

    def formula_correct_staticized(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Case_Decisions"]["I2"], "value", 540))

    def formula_status_staticized(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Case_Decisions"]["M3"], "value", "ESCALATED"))

    def formula_hold_staticized(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Case_Decisions"]["M4"], "value", "ON_HOLD"))

    def summary_staticized(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Control_Summary"]["B6"], "value", 7900))

    def routing_staticized(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: setattr(wb["Routing_Queue"]["E3"], "value", "payroll-director@company.example.com"))

    def decision_cell(cell: str, value):
        return lambda root: save_workbook(root / XLSX, lambda wb: setattr(wb["Case_Decisions"][cell], "value", value))

    def summary_cell(cell: str, value):
        return lambda root: save_workbook(root / XLSX, lambda wb: setattr(wb["Control_Summary"][cell], "value", value))

    def chart_removed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: wb["Control_Summary"]._charts.clear())

    def table_removed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: wb["Case_Decisions"].tables.clear())

    def conditional_formatting_removed(root: Path) -> None:
        save_workbook(root / XLSX, lambda wb: wb["Case_Decisions"].conditional_formatting._cf_rules.clear())

    def mbox_drop_director(root: Path) -> None:
        mutate_mbox(root / MBOX, lambda messages: messages.pop(4))

    def mbox_drop_elise(root: Path) -> None:
        mutate_mbox(root / MBOX, lambda messages: messages.pop(0))

    def mbox_add_hannah_recipient(root: Path) -> None:
        mutate_mbox(root / MBOX, lambda messages: messages[0].replace_header("To", "hannah.olsen@company.example.com"))

    def mbox_manager_recipient(root: Path) -> None:
        mutate_mbox(root / MBOX, lambda messages: messages[1].replace_header("To", "team-lead-b@company.example.com"))

    def mbox_director_drop_george(root: Path) -> None:
        mutate_mbox(root / MBOX, lambda messages: replace_message_body(messages[4], "George Nakamura", "Missing employee"))

    def mbox_director_add_hannah(root: Path) -> None:
        mutate_mbox(root / MBOX, lambda messages: replace_message_body(messages[4], "One disputed $150 case", "Hannah Olsen disputed $150 case"))

    def mbox_wrong_elise_recipient(root: Path) -> None:
        mutate_mbox(root / MBOX, lambda messages: messages[0].replace_header("To", "george.nakamura@company.example.com"))

    def mbox_false_sent(root: Path) -> None:
        mutate_mbox(root / MBOX, lambda messages: messages[0].replace_header("X-EvoCowork-Status", "SENT"))

    def mbox_claim_executed(root: Path) -> None:
        mutate_mbox(root / MBOX, lambda messages: replace_message_body(messages[0], "No payment", "Back-pay issued; payment completed. No payment"))

    def mbox_cross_employee(root: Path) -> None:
        mutate_mbox(root / MBOX, lambda messages: replace_message_body(messages[0], "Elise Fontaine,", "Elise Fontaine and Kevin Obi,"))

    def slack_posted(root: Path) -> None:
        mutate_json(root / SLACK, lambda payload: payload.update({"status": "POSTED", "posted": True}))

    def slack_channel_wrong(root: Path) -> None:
        mutate_json(root / SLACK, lambda payload: payload["source_channel"].update({"id": "C_WRONG"}))

    def slack_count_wrong(root: Path) -> None:
        mutate_json(root / SLACK, lambda payload: payload["summary"].update({"resolved": 5}))

    def slack_amount_wrong(root: Path) -> None:
        mutate_json(root / SLACK, lambda payload: payload["summary"].update({"actionable_absolute_variance_usd": 7900}))

    def slack_leak_name(root: Path) -> None:
        mutate_json(root / SLACK, lambda payload: payload["draft_message"].update({"text": payload["draft_message"]["text"] + " Elise Fontaine"}))

    def slack_leak_email(root: Path) -> None:
        mutate_json(root / SLACK, lambda payload: payload.update({"owner": "elise.fontaine@company.example.com"}))

    def slack_leak_amount(root: Path) -> None:
        mutate_json(root / SLACK, lambda payload: payload["draft_message"].update({"text": payload["draft_message"]["text"] + " $180"}))

    def slack_trace_missing(root: Path) -> None:
        mutate_json(root / SLACK, lambda payload: payload.update({"source_trace": []}))

    def control_executed(root: Path) -> None:
        mutate_json(root / CONTROL, lambda payload: payload.update({"decision": "EXECUTED", "status": "EXECUTED", "executed": True, "execution_record_count": 7}))

    def control_transaction(root: Path) -> None:
        mutate_json(root / CONTROL, lambda payload: payload.update({"transaction_ids": ["TX-FAKE-1"]}))

    def control_source_row(root: Path) -> None:
        mutate_json(root / CONTROL, lambda payload: payload["source_policy"].update({"source_row": 4}))

    def control_decision_wrong(root: Path) -> None:
        mutate_json(root / CONTROL, lambda payload: payload.update({"decision": "APPROVED_TO_EXECUTE"}))

    def control_trace_missing(root: Path) -> None:
        mutate_json(root / CONTROL, lambda payload: payload.update({"source_trace": []}))

    def word_total_wrong(root: Path) -> None:
        replace_docx(root / DOCX, "$7,900", "$7,800")

    def word_hannah_wrong(root: Path) -> None:
        replace_docx(root / DOCX, "ON_HOLD\nNo email", "RESOLVED\nEmployee draft")

    def word_george_amount_wrong(root: Path) -> None:
        replace_docx(root / DOCX, "$2,000 already paid; current correct $0", "$2,000 unpaid; current correct $2,000")

    def word_false_execution(root: Path) -> None:
        replace_docx(root / DOCX, "Nothing was sent or posted", "Email was sent and Slack was posted")

    def word_policy_removed(root: Path) -> None:
        replace_docx(root / DOCX, "HR Ops must route — never execute — back-pay corrections.", "HR Ops may execute back-pay corrections.")

    return {
        "missing-docx": missing_docx,
        "extra-debug-file": extra_debug,
        "corrupt-xlsx": corrupt_xlsx,
        "source-name-changed": source_name_changed,
        "source-issue-changed": source_issue_changed,
        "source-recorded-changed": source_recorded_changed,
        "source-note-changed": source_note_changed,
        "policy-threshold-changed": policy_threshold_changed,
        "policy-backpay-changed": policy_backpay_changed,
        "evidence-mail-removed": evidence_mail_removed,
        "evidence-elise-amount-wrong": evidence_elise_amount,
        "evidence-george-amount-wrong": evidence_george_amount,
        "evidence-channel-wrong": evidence_channel_wrong,
        "evidence-time-wrong": evidence_time_wrong,
        "formula-correct-staticized": formula_correct_staticized,
        "formula-status-staticized": formula_status_staticized,
        "formula-hold-staticized": formula_hold_staticized,
        "summary-staticized": summary_staticized,
        "routing-staticized": routing_staticized,
        "elise-correct-wrong": decision_cell("I2", 360),
        "george-correct-wrong": decision_cell("I3", 2000),
        "george-marked-resolved": decision_cell("M3", "RESOLVED"),
        "hannah-marked-resolved": decision_cell("M4", "RESOLVED"),
        "hannah-marked-escalated": decision_cell("M4", "ESCALATED"),
        "ibrahim-correct-wrong": decision_cell("I5", 3200),
        "ibrahim-marked-escalated": decision_cell("M5", "ESCALATED"),
        "jenny-marked-resolved": decision_cell("M6", "RESOLVED"),
        "kevin-marked-escalated": decision_cell("M7", "ESCALATED"),
        "lina-marked-escalated": decision_cell("M8", "ESCALATED"),
        "miguel-marked-resolved": decision_cell("M9", "RESOLVED"),
        "george-signed-positive": decision_cell("J3", 2000),
        "ibrahim-signed-positive": decision_cell("J5", 350),
        "george-absolute-negative": decision_cell("K3", -2000),
        "summary-total-wrong": summary_cell("B6", 7800),
        "summary-actionable-wrong": summary_cell("B7", 7900),
        "summary-draft-count-wrong": summary_cell("B9", 7),
        "native-chart-removed": chart_removed,
        "native-table-removed": table_removed,
        "conditional-formatting-removed": conditional_formatting_removed,
        "mbox-drop-director": mbox_drop_director,
        "mbox-drop-elise": mbox_drop_elise,
        "mbox-add-hannah-recipient": mbox_add_hannah_recipient,
        "mbox-manager-recipient": mbox_manager_recipient,
        "mbox-director-drop-george": mbox_director_drop_george,
        "mbox-director-add-hannah": mbox_director_add_hannah,
        "mbox-wrong-elise-recipient": mbox_wrong_elise_recipient,
        "mbox-false-sent": mbox_false_sent,
        "mbox-claim-executed": mbox_claim_executed,
        "mbox-cross-employee": mbox_cross_employee,
        "slack-false-posted": slack_posted,
        "slack-wrong-channel": slack_channel_wrong,
        "slack-count-wrong": slack_count_wrong,
        "slack-amount-wrong": slack_amount_wrong,
        "slack-leak-name": slack_leak_name,
        "slack-leak-email": slack_leak_email,
        "slack-leak-case-amount": slack_leak_amount,
        "slack-source-trace-missing": slack_trace_missing,
        "control-false-executed": control_executed,
        "control-fake-transaction": control_transaction,
        "control-source-row-wrong": control_source_row,
        "control-decision-wrong": control_decision_wrong,
        "control-source-trace-missing": control_trace_missing,
        "word-total-wrong": word_total_wrong,
        "word-hannah-wrong": word_hannah_wrong,
        "word-george-amount-wrong": word_george_amount_wrong,
        "word-false-execution": word_false_execution,
        "word-policy-removed": word_policy_removed,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--golden", type=Path, default=Path(__file__).resolve().parents[1] / "solution" / "Golden")
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    judge_path = Path(__file__).with_name("judge.py")
    spec = importlib.util.spec_from_file_location("payroll_control_task_judge", judge_path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入本题 Judge")
    judge = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = judge
    spec.loader.exec_module(judge)

    results = []
    with tempfile.TemporaryDirectory(prefix="payroll-control-negative-cases-") as temporary:
        root = Path(temporary)
        for name, mutate in cases().items():
            case_dir = root / name
            shutil.copytree(args.golden, case_dir)
            mutate(case_dir)
            result = judge.evaluate(case_dir)
            results.append({
                "case": name,
                "reward": result["reward"],
                "passed": result["passed"],
                "hurdle_failures": result.get("hurdle_failures", []),
                "penalties": [item["id"] for item in result.get("penalties", [])],
            })
    failed = [item for item in results if item["passed"] or float(item["reward"]) >= 0.8]
    report = {
        "total": len(results),
        "rejected": len(results) - len(failed),
        "maximum_reward": max(float(item["reward"]) for item in results),
        "cases": results,
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if failed:
        raise SystemExit(f"存在误接受负例：{[item['case'] for item in failed]}")


if __name__ == "__main__":
    main()
