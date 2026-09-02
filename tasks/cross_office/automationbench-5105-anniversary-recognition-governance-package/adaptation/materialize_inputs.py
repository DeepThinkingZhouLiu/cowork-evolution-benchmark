#!/usr/bin/env python3
"""只静态化 AutomationBench 5105；不是通用造题器。"""
import hashlib,importlib.util,json,shutil,sys,types
from datetime import date,datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.worksheet.table import Table,TableStyleInfo
ROOT=Path(__file__).resolve().parents[1];UP=Path(__file__).with_name("upstream-hr-tasks.py");SNAP=Path(__file__).with_name("upstream-task-5105.json");DATA=ROOT/"data/input_files";ENV=ROOT/"environment/input_files";UPSTREAM_SHA="c173a36a03f43096bd122577584aa32f9df379d8760aa46c43be04c85d101779"
def load_task():
 if hashlib.sha256(UP.read_bytes()).hexdigest()!=UPSTREAM_SHA:raise ValueError("上游哈希错误")
 d=types.ModuleType("datasets");d.Dataset=object;sys.modules["datasets"]=d
 for name in ("automationbench","automationbench.domains","automationbench.domains.hr"):sys.modules[name]=types.ModuleType(name)
 noise=types.ModuleType("automationbench.domains.hr._noise");noise.apply_noise=lambda x:x;sys.modules[noise.__name__]=noise
 spec=importlib.util.spec_from_file_location("automationbench_5105",UP);module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module);task=module.get_task_5105();assert task["example_id"]==5105 and task["task"]=="hr.anniversary_recognition";return task
def audit(task):
 info=task["info"];state=info["initial_state"];assert state["meta"]["current_time"]=="2026-04-15T09:00:00Z" and len(info["assertions"])==9
 rows=state["google_sheets"]["spreadsheets"][0]["worksheets"][0]["rows"];assert [r["cells"]["Name"] for r in rows]==["Angela Reeves","Ben Harrington","Celine Dubois","Derek Huang","Eva Lindström","Frank Torres","Grace Okonkwo","Henry Zhao","Irene Kowalski","Jake Morrison"]
 messages=state["gmail"]["messages"];assert [m["id"] for m in messages]==["msg_anniversary_last_5105","msg_exec_cc_policy_5105","msg_hire_correction_5105"] and "do NOT CC" in messages[1]["body_plain"] and "2024-04-01" in messages[2]["body_plain"]
def write_mbox(messages,path):
 chunks=[]
 for record in messages:
  msg=EmailMessage(policy=SMTP);msg["From"]=record["from_"];msg["To"]=", ".join(record["to"]);msg["Subject"]=record["subject"];dt=datetime.fromisoformat(record["date"].replace("Z","+00:00"));msg["Date"]=dt.strftime("%a, %d %b %Y %H:%M:%S +0000");msg["Message-ID"]=f"<{record['id']}@automationbench.invalid>";msg["X-AutomationBench-Message-ID"]=record["id"];msg.set_content(record["body_plain"],charset="utf-8");buf=BytesIO();BytesGenerator(buf,policy=SMTP,mangle_from_=True,maxheaderlen=0).flatten(msg);chunks.extend([f"From {record['from_']} {dt.strftime('%a %b %d %H:%M:%S %Y')}\n".encode(),buf.getvalue().replace(b"\r\n",b"\n"),b"\n"])
 path.write_bytes(b"".join(chunks))
def write_xlsx(state,path):
 rows=state["google_sheets"]["spreadsheets"][0]["worksheets"][0]["rows"];headers=["Source Row","Name","Department","Hire Date Raw","Status"];wb=Workbook();ws=wb.active;ws.title="Employees";ws.append(headers)
 for row in rows:
  c=row["cells"];raw=c["Hire Date"]
  try:value=date.fromisoformat(raw)
  except (ValueError,TypeError):value=raw
  ws.append([row["row_id"],c["Name"],c["Department"],value,c["Status"]])
 for i,w in enumerate([10,24,22,20,16],1):ws.column_dimensions[ws.cell(1,i).column_letter].width=w
 ws.freeze_panes="A2";ws.sheet_view.showGridLines=False
 for cell in ws[1]:cell.fill=PatternFill("solid",fgColor="17365D");cell.font=Font(name="Carlito",bold=True,color="FFFFFF");cell.alignment=Alignment(horizontal="center",wrap_text=True)
 for row in ws.iter_rows(min_row=2):
  for cell in row:cell.font=Font(name="Carlito",size=9);cell.alignment=Alignment(vertical="top",wrap_text=True)
  if isinstance(row[3].value,date):row[3].number_format="yyyy-mm-dd"
 table=Table(displayName="EmployeeDirectorySource",ref=ws.dimensions);table.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True);ws.add_table(table);ws.page_setup.orientation="landscape";ws.page_setup.fitToWidth=1;ws.sheet_properties.pageSetUpPr.fitToPage=True;wb.save(path)
def main():
 task=load_task();audit(task);state=task["info"]["initial_state"];DATA.mkdir(parents=True,exist_ok=True);ENV.mkdir(parents=True,exist_ok=True);snapshot=json.dumps(task,ensure_ascii=False,indent=2)+"\n"
 if SNAP.exists() and SNAP.read_text()!=snapshot:raise ValueError("单题快照漂移")
 SNAP.write_text(snapshot);write_xlsx(state,DATA/"Employee_Directory.xlsx");write_mbox(state["gmail"]["messages"],DATA/"Anniversary_Policies_and_Correction.mbox");(DATA/"General_Slack.json").write_text(json.dumps(state["slack"],ensure_ascii=False,indent=2)+"\n");(DATA/"Workflow_Context.json").write_text(json.dumps({"meta":state["meta"]},indent=2)+"\n")
 for item in DATA.iterdir():
  if item.is_file():shutil.copy2(item,ENV/item.name)
if __name__=="__main__":main()
