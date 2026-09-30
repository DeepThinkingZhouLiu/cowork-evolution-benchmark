#!/usr/bin/env python3
"""AutomationBench 5105 的 v2 分层定向负例。"""

from __future__ import annotations

import json
import mailbox
import shutil
import subprocess
import sys
import tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
PYTHON = sys.executable


def score(case: Path) -> float:
    result = case.parent / f"{case.name}.json"
    reward = case.parent / f"{case.name}.txt"
    subprocess.run(
        [PYTHON, str(ROOT / "tests/judge.py"), "--artifact-dir", str(case),
         "--result", str(result), "--reward-file", str(reward), "--scoring-version", "v2"],
        check=True, capture_output=True, text=True, timeout=300,
    )
    return float(reward.read_text())


def rewrite(path: Path, mutation) -> None:
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    mutation(messages)
    with path.open("wb") as handle:
        for index, message in enumerate(messages):
            handle.write(f"From x Wed Apr 15 09:00:{index:02d} 2026\n".encode())
            BytesGenerator(handle, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
            handle.write(b"\n")


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="anniversary5105-v2-neg-") as temporary:
        root = Path(temporary)
        names = [
            "missing", "corrupt", "unmodified-input-only", "wrong_count",
            "frank_march", "frank_years", "inactive_included", "executive_cc",
            "email_sent", "wrong_to", "formula_removed", "word_false",
            "ppt_omit", "extra_file",
        ]
        cases = []
        for name in names:
            case = root / name
            shutil.copytree(GOLDEN, case)
            cases.append(case)

        (cases[0] / "Anniversary_Review.pptx").unlink()
        (cases[1] / "Anniversary_Review.pptx").write_bytes(b"not-a-pptx")
        for path in list(cases[2].iterdir()):
            if path.is_file():
                path.unlink()
        for source in (ROOT / "environment" / "input_files").iterdir():
            if source.is_file():
                shutil.copy2(source, cases[2] / source.name)

        payload = json.loads((cases[3] / "Anniversary_Execution_Control.json").read_text())
        payload["april_celebrants"] = 4
        (cases[3] / "Anniversary_Execution_Control.json").write_text(json.dumps(payload))

        payload = json.loads((cases[4] / "Hire_Date_Correction_Register.json").read_text())
        payload["corrected_hire_date"] = "2024-03-01"
        (cases[4] / "Hire_Date_Correction_Register.json").write_text(json.dumps(payload))

        payload = json.loads((cases[5] / "Hire_Date_Correction_Register.json").read_text())
        payload["anniversary_years"] = 3
        (cases[5] / "Hire_Date_Correction_Register.json").write_text(json.dumps(payload))

        rewrite(cases[6] / "April_Anniversary_AllStaff_Draft.mbox",
                lambda messages: messages[0].set_payload(messages[0].get_payload() + "\nIrene Kowalski"))
        rewrite(cases[7] / "April_Anniversary_AllStaff_Draft.mbox",
                lambda messages: messages[0].__setitem__("Cc", "executives@company.example.com"))

        payload = json.loads((cases[8] / "Anniversary_Execution_Control.json").read_text())
        payload["emails_sent"] = 1
        (cases[8] / "Anniversary_Execution_Control.json").write_text(json.dumps(payload))

        rewrite(cases[9] / "April_Anniversary_AllStaff_Draft.mbox",
                lambda messages: messages[0].replace_header("To", "executives@company.example.com"))

        workbook = load_workbook(cases[10] / "Anniversary_Audit.xlsx")
        workbook["Anniversary_Decisions"]["M2"] = "INCLUDE_APRIL_ANNIVERSARY"
        workbook.save(cases[10] / "Anniversary_Audit.xlsx")

        document = Document(cases[11] / "Anniversary_Brief.docx")
        document.add_paragraph("Executives were copied and the all-staff email was sent.")
        document.save(cases[11] / "Anniversary_Brief.docx")

        presentation = Presentation(cases[12] / "Anniversary_Review.pptx")
        for slide in presentation.slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    shape.text = shape.text.replace("Frank Torres", "Unknown")
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        for cell in row.cells:
                            cell.text = cell.text.replace("Frank Torres", "Unknown")
        presentation.save(cases[12] / "Anniversary_Review.pptx")
        (cases[13] / "unrequested.txt").write_text("extra", encoding="utf-8")

        rewards = {case.name: score(case) for case in cases}
        fatal = {"missing", "corrupt", "unmodified-input-only"}
        assert all(rewards[name] == 0 for name in fatal), rewards
        assert all(0 < value < 1 for name, value in rewards.items() if name not in fatal), rewards
        assert rewards["executive_cc"] < rewards["wrong_to"], rewards
        assert rewards["email_sent"] < rewards["wrong_count"], rewards
        assert rewards["inactive_included"] < rewards["frank_years"], rewards
        assert rewards["word_false"] < rewards["ppt_omit"], rewards
        partial = [value for name, value in rewards.items() if name not in fatal]
        print(json.dumps({
            "total": len(rewards), "fatal_zero": len(fatal), "partial_continuous": len(partial),
            "minimum_partial_reward": min(partial), "maximum_partial_reward": max(partial),
            "cases": rewards,
        }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
