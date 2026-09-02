#!/usr/bin/env python3
"""AutomationBench 5106 培训合规执法包的逐题参考解法。"""

from __future__ import annotations

import argparse
import hashlib
import json
import mailbox
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


INPUT_HASHES = {
    "Training_Compliance_Tracker.xlsx": "5243ebc360d34f8ff167b91a5ffa8b0d437ea31fd82b65129bfe637573e33394",
    "Compliance_Handbook.mbox": "dd2f164ce3d3370a58e03f657eb423c0ecb80cc637b371887af6cd46cf05d592",
    "Compliance_Channel.json": "947f627f9b48266a49adecce23590af7ba9d516984d3851796b346709fb8f169",
    "Workflow_Context.json": "76c4a0657bc27f2fd830b10613ec7ceae3adcdc5798aeee0f822d27bf69bbc99",
}
XLSX_NAME = "Training_Enforcement_Audit.xlsx"
DOCX_NAME = "Compliance_Enforcement_Brief.docx"
MBOX_NAME = "Enforcement_Notification_Drafts.mbox"
SLACK_NAME = "Compliance_Channel_Draft.json"
CONTROL_NAME = "Access_Suspension_Control.json"
FIXED_ISO_TIME = "2026-04-15T09:00:00Z"
FIXED_DATE = date(2026, 4, 15)
FIXED_ZIP_TIME = (2026, 4, 15, 9, 0, 0)

NAVY = "17365D"
GREEN = "E2F0D9"
AMBER = "FFF2CC"
RED = "F4CCCC"
PURPLE = "E4DFEC"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="B7C9D6")


@dataclass(frozen=True)
class TrainingCase:
    source_row: int
    employee: str
    employee_email: str
    manager: str
    manager_email: str
    training: str
    due_date: str
    source_status: str
    notes: str


@dataclass(frozen=True)
class Decision:
    source: TrainingCase
    days_overdue: int
    completion_override: bool
    active_exemption: bool
    action: str
    primary_route: str
    secondary_route: str
    notification_plan: str


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def decode_message(message) -> str:
    if message.is_multipart():
        parts: list[str] = []
        for part in message.walk():
            if part.get_content_type() == "text/plain":
                payload = part.get_payload(decode=True) or b""
                parts.append(payload.decode(part.get_content_charset() or "utf-8", errors="replace"))
        return "\n".join(parts).strip()
    payload = message.get_payload(decode=True) or b""
    return payload.decode(message.get_content_charset() or "utf-8", errors="replace").strip()


def read_inputs(input_dir: Path) -> tuple[list[TrainingCase], dict, dict, str]:
    for name, expected in INPUT_HASHES.items():
        source = input_dir / name
        if not source.is_file() or sha256(source) != expected:
            raise ValueError(f"输入 {name} 缺失或 SHA256 与冻结版本不一致")

    context = json.loads((input_dir / "Workflow_Context.json").read_text(encoding="utf-8"))
    if context != {"meta": {"schema_version": "0.1.0", "current_time": FIXED_ISO_TIME}}:
        raise ValueError("冻结运行时间已变化")

    channel = json.loads((input_dir / "Compliance_Channel.json").read_text(encoding="utf-8"))
    expected_channel = {
        "channels": [{"id": "C_COMP_5106", "name": "compliance", "is_private": False}],
        "messages": [],
    }
    if channel != expected_channel:
        raise ValueError("compliance Slack 频道快照已变化")

    workbook = load_workbook(input_dir / "Training_Compliance_Tracker.xlsx", read_only=True, data_only=True)
    if workbook.sheetnames != ["Training Status"]:
        raise ValueError("培训合规工作簿结构已变化")
    cases: list[TrainingCase] = []
    for row in workbook["Training Status"].iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        cases.append(TrainingCase(
            source_row=int(row[0]), employee=str(row[1]), employee_email=str(row[2]),
            manager=str(row[3]), manager_email=str(row[4]), training=str(row[5]),
            due_date=str(row[6]), source_status=str(row[7]), notes=str(row[8] or ""),
        ))
    workbook.close()
    frozen = [
        (2, "Patricia Chen", "patricia.chen@company.example.com", "Vikram Shah", "vikram.shah@company.example.com", "Information Security", "2026-04-01", "Overdue", ""),
        (3, "Patricia Chang", "patricia.chang@company.example.com", "Vikram Shah", "vikram.shah@company.example.com", "Anti-Harassment", "2026-03-10", "Overdue", ""),
        (4, "Roberto Alvarez", "roberto.alvarez@company.example.com", "Diana Torres", "diana.torres@company.example.com", "Data Privacy", "2026-03-01", "Overdue", "Medical leave exemption through April 30"),
        (5, "Samantha Green", "samantha.green@company.example.com", "Diana Torres", "diana.torres@company.example.com", "Information Security", "2026-03-25", "Overdue", ""),
        (6, "Theodore Kim", "theodore.kim@company.example.com", "Vikram Shah", "vikram.shah@company.example.com", "Anti-Harassment", "2026-02-10", "Overdue", ""),
        (7, "Ursula Martin", "ursula.martin@company.example.com", "Diana Torres", "diana.torres@company.example.com", "Data Privacy", "2026-04-10", "Overdue", ""),
        (8, "Victor Pham", "victor.pham@company.example.com", "Vikram Shah", "vikram.shah@company.example.com", "Information Security", "2025-12-15", "Overdue", ""),
        (9, "Wendy Torres", "wendy.torres@company.example.com", "Diana Torres", "diana.torres@company.example.com", "Anti-Harassment", "2026-04-14", "Complete", ""),
    ]
    actual = [
        (item.source_row, item.employee, item.employee_email, item.manager, item.manager_email,
         item.training, item.due_date, item.source_status, item.notes)
        for item in cases
    ]
    if actual != frozen:
        raise ValueError("八条冻结培训记录已变化")

    box = mailbox.mbox(input_dir / "Compliance_Handbook.mbox", create=False)
    messages = list(box)
    box.close()
    if len(messages) != 1:
        raise ValueError("手册邮箱必须恰好包含一封邮件")
    message = messages[0]
    message_id = str(message.get("X-AutomationBench-Message-ID", ""))
    body = decode_message(message)
    required = (
        "COMPLIANCE HANDBOOK 2026", "Section 6: Enforcement Actions",
        "1 to 14 days overdue", "15 to 30 days overdue", "Over 30 days overdue",
        "compliance-committee@company.example.com", "active exemption",
        "do not CC or otherwise notify it for 1-to-30-day",
    )
    if message_id != "msg_handbook_5106" or any(token not in body for token in required):
        raise ValueError("冻结合规手册证据不完整")
    handbook = {
        "message_id": message_id,
        "thread_id": str(message.get("X-AutomationBench-Thread-ID", "")),
        "from": str(message.get("From", "")),
        "to": str(message.get("To", "")),
        "subject": str(message.get("Subject", "")),
        "date": str(message.get("Date", "")),
        "body": body,
    }
    return cases, handbook, channel, context["meta"]["current_time"]


def derive_decisions(cases: list[TrainingCase]) -> list[Decision]:
    decisions: list[Decision] = []
    for item in cases:
        due = date.fromisoformat(item.due_date)
        days_overdue = max(0, (FIXED_DATE - due).days)
        completion = item.source_status.casefold() == "complete"
        exemption = item.employee == "Roberto Alvarez" and item.notes == "Medical leave exemption through April 30"
        if completion:
            action = "NO_ACTION"
        elif exemption:
            action = "EXEMPT"
        elif days_overdue == 0:
            action = "NO_ACTION"
        elif days_overdue <= 14:
            action = "WARNING"
        elif days_overdue <= 30:
            action = "SUSPENSION_REQUEST"
        else:
            action = "COMMITTEE_ESCALATION"

        if action == "WARNING":
            primary, secondary, plan = item.employee_email, item.manager_email, "EMPLOYEE_WARNING_CC_MANAGER"
        elif action == "SUSPENSION_REQUEST":
            primary, secondary, plan = "it-security@company.example.com", item.manager_email, "IT_REQUEST_AND_MANAGER_NOTICE"
        elif action == "COMMITTEE_ESCALATION":
            primary, secondary, plan = "compliance-committee@company.example.com", "NONE", "COMMITTEE_ONLY"
        else:
            primary, secondary, plan = "NONE", "NONE", "NO_DRAFT"
        decisions.append(Decision(item, days_overdue, completion, exemption, action, primary, secondary, plan))

    expected = {
        "Patricia Chen": (14, "WARNING"),
        "Patricia Chang": (36, "COMMITTEE_ESCALATION"),
        "Roberto Alvarez": (45, "EXEMPT"),
        "Samantha Green": (21, "SUSPENSION_REQUEST"),
        "Theodore Kim": (64, "COMMITTEE_ESCALATION"),
        "Ursula Martin": (5, "WARNING"),
        "Victor Pham": (121, "COMMITTEE_ESCALATION"),
        "Wendy Torres": (1, "NO_ACTION"),
    }
    if {item.source.employee: (item.days_overdue, item.action) for item in decisions} != expected:
        raise ValueError("逐案执法真值与冻结任务不一致")
    return decisions


def add_table(sheet, name: str) -> None:
    table = Table(displayName=name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)


def style_sheet(sheet, widths: dict[str, float], footer: str, *, a3: bool = False) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", bold=True, color=WHITE, size=9)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)
    sheet.row_dimensions[1].height = 30
    grid = Side(style="hair", color="D9E2F3")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=8.5)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=grid, right=grid, bottom=grid)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if a3 else sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.oddFooter.center.text = footer
    sheet.print_area = sheet.dimensions


def normalize_ooxml(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="compliance-5106-ooxml-") as temporary:
        target = Path(temporary) / path.name
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
            target, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
        ) as destination:
            axis_ids: dict[str, str] = {}
            for name in sorted(source.namelist()):
                data = source.read(name)
                if name == "docProps/core.xml":
                    try:
                        root = etree.fromstring(data)
                        namespaces = {
                            "dcterms": "http://purl.org/dc/terms/",
                            "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
                        }
                        for tag in ("created", "modified"):
                            for node in root.xpath(f"//dcterms:{tag}", namespaces=namespaces):
                                node.text = "2026-04-15T09:00:00Z"
                        for tag in ("lastModifiedBy", "revision"):
                            for node in root.xpath(f"//cp:{tag}", namespaces=namespaces):
                                node.text = "EvoCowork" if tag == "lastModifiedBy" else "1"
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except etree.XMLSyntaxError:
                        pass
                if name.startswith("xl/charts/chart") and name.endswith(".xml"):
                    try:
                        root = etree.fromstring(data)
                        namespaces = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
                        for node in root.xpath("//c:axId", namespaces=namespaces):
                            old = node.get("val")
                            if old and old not in axis_ids:
                                axis_ids[old] = str(51_060_001 + len(axis_ids))
                        for node in root.xpath("//c:axId | //c:crossAx", namespaces=namespaces):
                            old = node.get("val")
                            if old in axis_ids:
                                node.set("val", axis_ids[old])
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except etree.XMLSyntaxError:
                        pass
                info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                destination.writestr(info, data)
        shutil.copy2(target, path)


def recalculate_xlsx(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="compliance-5106-recalc-") as temporary:
        root = Path(temporary)
        source_dir, output_dir, profile_dir = root / "source", root / "output", root / "profile"
        source_dir.mkdir()
        output_dir.mkdir()
        source = source_dir / path.name
        shutil.copy2(path, source)
        process = subprocess.run(
            [
                "libreoffice", "--headless", f"-env:UserInstallation=file://{profile_dir}",
                "--convert-to", "xlsx", "--outdir", str(output_dir), str(source),
            ],
            stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300,
        )
        converted = output_dir / path.name
        if process.returncode != 0 or not converted.is_file():
            raise RuntimeError(f"LibreOffice 公式重算失败：{process.stdout[-1000:]}")
        shutil.copy2(converted, path)


def write_workbook(
    path: Path, cases: list[TrainingCase], decisions: list[Decision], handbook: dict,
    channel: dict, frozen_time: str,
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.creator = "EvoCowork AutomationBench 5106 adaptation"
    workbook.title = "2026 Compliance Training Enforcement Audit"
    workbook.subject = "Formula-driven policy and routing audit"
    workbook.description = "All people, dates, statuses, exemptions and rules trace to frozen AutomationBench 5106 inputs."
    fixed_dt = datetime(2026, 4, 15, 9, 0, 0)
    workbook.created = fixed_dt
    workbook.modified = fixed_dt

    source = workbook.create_sheet("Source_Training")
    source.append([
        "Source Row", "Employee", "Employee Email", "Manager", "Manager Email",
        "Training", "Due Date", "Status", "Notes",
    ])
    for item in cases:
        source.append([
            item.source_row, item.employee, item.employee_email, item.manager, item.manager_email,
            item.training, item.due_date, item.source_status, item.notes,
        ])
    add_table(source, "ComplianceSourceTraining")
    style_sheet(source, {
        "A": 11, "B": 21, "C": 32, "D": 20, "E": 32,
        "F": 25, "G": 15, "H": 13, "I": 48,
    }, "Frozen source rows · AutomationBench 5106", a3=True)
    for row in range(2, 10):
        source.row_dimensions[row].height = 30
    source.row_dimensions[4].height = 42

    policy = workbook.create_sheet("Policy_Rules")
    policy.append(["Section", "Rule", "Required Action", "Restriction", "Source Trace"])
    policy_rows = [
        ("Section 6", "1 to 14 days overdue", "Warning email to employee; CC manager", "Do not notify compliance committee", "Compliance_Handbook.mbox:msg_handbook_5106"),
        ("Section 6", "15 to 30 days overdue", "Request IT access suspension; email manager", "Do not notify compliance committee", "Compliance_Handbook.mbox:msg_handbook_5106"),
        ("Section 6", "Over 30 days overdue", "Escalate to compliance committee", "Do NOT take direct action", "Compliance_Handbook.mbox:msg_handbook_5106"),
        ("Section 6", "Active exemption in Notes", "Exclude from enforcement", "Overrides overdue status and days", "Compliance_Handbook.mbox:msg_handbook_5106"),
    ]
    for row in policy_rows:
        policy.append(row)
    add_table(policy, "CompliancePolicyRules")
    style_sheet(policy, {"A": 16, "B": 28, "C": 52, "D": 48, "E": 58}, "Section 6 · frozen enforcement policy", a3=True)
    for row in range(2, 6):
        policy.row_dimensions[row].height = 46

    evidence = workbook.create_sheet("Evidence_Register")
    evidence.append(["Record", "Evidence Type", "Evidence Ref", "Evidence Detail", "Frozen At", "Source Trace"])
    for item in cases:
        evidence.append([
            f"Source row {item.source_row} · {item.employee}", "SOURCE_ROW", f"Training Status row {item.source_row}",
            f"{item.training}; due {item.due_date}; status {item.source_status}; notes {item.notes or '[blank]'}",
            frozen_time, f"Training_Compliance_Tracker.xlsx:Training Status row {item.source_row}",
        ])
    evidence.append([
        "Compliance handbook", "FULL_EMAIL", handbook["message_id"], handbook["body"], handbook["date"],
        "Compliance_Handbook.mbox:msg_handbook_5106",
    ])
    evidence.append(["Workflow context", "FROZEN_TIME", "meta.current_time", frozen_time, frozen_time, "Workflow_Context.json:meta.current_time"])
    evidence.append([
        "Compliance channel", "CHANNEL_SNAPSHOT", channel["channels"][0]["id"],
        json.dumps(channel["channels"][0], ensure_ascii=False, sort_keys=True), frozen_time,
        "Compliance_Channel.json:channels[0]",
    ])
    add_table(evidence, "ComplianceEvidenceRegister")
    style_sheet(evidence, {"A": 32, "B": 23, "C": 34, "D": 92, "E": 24, "F": 68}, "Evidence register · source rows, handbook, time and channel", a3=True)
    for row in range(2, 10):
        evidence.row_dimensions[row].height = 34
    evidence.row_dimensions[10].height = 170
    evidence.row_dimensions[11].height = 28
    evidence.row_dimensions[12].height = 34

    decisions_sheet = workbook.create_sheet("Enforcement_Decisions")
    decisions_sheet.append([
        "Source Row", "Employee", "Employee Email", "Manager", "Manager Email", "Training",
        "Due Date", "Source Status", "Notes", "Days Overdue", "Completion Override",
        "Active Exemption", "Final Action", "Primary Route", "Secondary Route",
        "Notification Plan", "Source Trace",
    ])
    for row in range(2, 10):
        decisions_sheet.append([
            f"=Source_Training!A{row}", f"=Source_Training!B{row}", f"=Source_Training!C{row}",
            f"=Source_Training!D{row}", f"=Source_Training!E{row}", f"=Source_Training!F{row}",
            f"=Source_Training!G{row}", f"=Source_Training!H{row}", f"=Source_Training!I{row}",
            f"=MAX(0,DATE(2026,4,15)-DATEVALUE(G{row}))",
            f'=IF(H{row}="Complete","YES","NO")',
            f'=IF(AND(ISNUMBER(SEARCH("exemption",I{row})),DATE(2026,4,15)<=DATE(2026,4,30)),"YES","NO")',
            f'=IF(K{row}="YES","NO_ACTION",IF(L{row}="YES","EXEMPT",IF(J{row}=0,"NO_ACTION",IF(J{row}<=14,"WARNING",IF(J{row}<=30,"SUSPENSION_REQUEST","COMMITTEE_ESCALATION")))))',
            f'=IF(M{row}="WARNING",C{row},IF(M{row}="SUSPENSION_REQUEST","it-security@company.example.com",IF(M{row}="COMMITTEE_ESCALATION","compliance-committee@company.example.com","NONE")))',
            f'=IF(OR(M{row}="WARNING",M{row}="SUSPENSION_REQUEST"),E{row},"NONE")',
            f'=IF(M{row}="WARNING","EMPLOYEE_WARNING_CC_MANAGER",IF(M{row}="SUSPENSION_REQUEST","IT_REQUEST_AND_MANAGER_NOTICE",IF(M{row}="COMMITTEE_ESCALATION","COMMITTEE_ONLY","NO_DRAFT")))',
            f'=Evidence_Register!F{row}',
        ])
    for action, color in (
        ("WARNING", AMBER), ("SUSPENSION_REQUEST", RED),
        ("COMMITTEE_ESCALATION", PURPLE), ("EXEMPT", GREEN), ("NO_ACTION", GRAY),
    ):
        decisions_sheet.conditional_formatting.add(
            "M2:M9", FormulaRule(formula=[f'M2="{action}"'], fill=PatternFill("solid", fgColor=color))
        )
    add_table(decisions_sheet, "ComplianceEnforcementDecisions")
    style_sheet(decisions_sheet, {
        "A": 10, "B": 20, "C": 30, "D": 18, "E": 30, "F": 23, "G": 14,
        "H": 13, "I": 45, "J": 14, "K": 18, "L": 17, "M": 25,
        "N": 38, "O": 34, "P": 34, "Q": 62,
    }, "Formula decisions · status/exemption precedence, then overdue bands", a3=True)
    for row in range(2, 10):
        decisions_sheet.row_dimensions[row].height = 38
    decisions_sheet.row_dimensions[4].height = 48

    queue = workbook.create_sheet("Notification_Queue")
    queue.append(["Draft ID", "Type", "To", "CC", "Included Cases", "Case Count", "Status", "Execution Boundary", "Source Trace"])
    queue_rows = [
        ("D001", "WARNING", "=Enforcement_Decisions!C2", "=Enforcement_Decisions!E2", "=Enforcement_Decisions!B2", 1, "DRAFT_NOT_SENT", "NO_SEND", "=Enforcement_Decisions!Q2"),
        ("D002", "WARNING", "=Enforcement_Decisions!C7", "=Enforcement_Decisions!E7", "=Enforcement_Decisions!B7", 1, "DRAFT_NOT_SENT", "NO_SEND", "=Enforcement_Decisions!Q7"),
        ("D003", "IT_SUSPENSION_REQUEST", "it-security@company.example.com", "", "=Enforcement_Decisions!B5", 1, "DRAFT_NOT_SENT", "REQUEST_ONLY_NO_EXECUTION", "=Enforcement_Decisions!Q5"),
        ("D004", "MANAGER_NOTICE", "=Enforcement_Decisions!E5", "", "=Enforcement_Decisions!B5", 1, "DRAFT_NOT_SENT", "NOTICE_ONLY", "=Enforcement_Decisions!Q5"),
        ("D005", "COMMITTEE_ESCALATION", "compliance-committee@company.example.com", "", '=Enforcement_Decisions!B3&", "&Enforcement_Decisions!B6&", "&Enforcement_Decisions!B8', 3, "DRAFT_NOT_SENT", "ESCALATE_ONLY_NO_DIRECT_ACTION", '=Enforcement_Decisions!Q3&"; "&Enforcement_Decisions!Q6&"; "&Enforcement_Decisions!Q8'),
    ]
    for row in queue_rows:
        queue.append(row)
    add_table(queue, "ComplianceNotificationQueue")
    style_sheet(queue, {"A": 12, "B": 27, "C": 40, "D": 36, "E": 58, "F": 13, "G": 20, "H": 34, "I": 72}, "Five offline drafts · no sends or access changes", a3=True)
    for row in range(2, 7):
        queue.row_dimensions[row].height = 42

    summary = workbook.create_sheet("Control_Summary")
    summary.append(["Metric", "Value", "Control Meaning"])
    summary_rows = [
        ("Total source cases", "=COUNTA(Enforcement_Decisions!B2:B9)", "All frozen source rows"),
        ("Warning", '=COUNTIF(Enforcement_Decisions!M2:M9,"WARNING")', "Employee warning + manager CC"),
        ("Suspension request", '=COUNTIF(Enforcement_Decisions!M2:M9,"SUSPENSION_REQUEST")', "IT request + manager notice; no execution"),
        ("Committee escalation", '=COUNTIF(Enforcement_Decisions!M2:M9,"COMMITTEE_ESCALATION")', "Committee only; no direct action"),
        ("Exempt", '=COUNTIF(Enforcement_Decisions!M2:M9,"EXEMPT")', "Active exemption overrides overdue"),
        ("No action", '=COUNTIF(Enforcement_Decisions!M2:M9,"NO_ACTION")', "Complete status overrides due date"),
        ("Enforcement cases", '=COUNTIF(Enforcement_Decisions!M2:M9,"WARNING")+COUNTIF(Enforcement_Decisions!M2:M9,"SUSPENSION_REQUEST")+COUNTIF(Enforcement_Decisions!M2:M9,"COMMITTEE_ESCALATION")', "Six routed cases"),
        ("Offline email drafts", "=COUNTA(Notification_Queue!A2:A6)", "Exactly five"),
        ("Slack messages posted", "=0", "Draft only"),
        ("Access suspensions executed", "=0", "Request draft only"),
    ]
    for row in summary_rows:
        summary.append(row)
    add_table(summary, "ComplianceControlSummary")
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Enforcement action distribution"
    chart.y_axis.title = "Cases"
    chart.x_axis.title = "Action"
    chart.height = 7.0
    chart.width = 14.0
    chart.add_data(Reference(summary, min_col=2, min_row=3, max_row=7), titles_from_data=False)
    chart.set_categories(Reference(summary, min_col=1, min_row=3, max_row=7))
    summary.add_chart(chart, "E2")
    style_sheet(summary, {"A": 32, "B": 18, "C": 52}, "Control totals · formulas linked to Enforcement_Decisions")
    summary.print_area = "A1:L11"

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(path)
    recalculate_xlsx(path)
    normalize_ooxml(path)


def set_cell_shading(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def format_doc_table(table, widths: list[float] | None = None, font_size: float = 8.0) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for cell in table.rows[0].cells:
        set_cell_shading(cell, NAVY)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(font_size)
    for row in table.rows:
        for col_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths and col_index < len(widths):
                cell.width = Cm(widths[col_index])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(font_size)


def add_title(document: Document, title: str, subtitle: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title)
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(23, 54, 93)
    subtitle_p = document.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle_p.add_run(subtitle)
    subtitle_run.italic = True
    subtitle_run.font.name = "Aptos"
    subtitle_run.font.size = Pt(10)
    subtitle_run.font.color.rgb = RGBColor(89, 89, 89)


def add_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(31, 78, 121)


def add_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(9.5)


def write_brief(path: Path, decisions: list[Decision]) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.3)
    section.bottom_margin = Cm(1.3)
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    core = document.core_properties
    core.title = "2026 Compliance Training Enforcement Brief"
    core.author = "EvoCowork AutomationBench 5106 adaptation"
    core.subject = "Eight-case compliance training policy and routing brief"
    core.created = datetime(2026, 4, 15, 9, 0, 0)
    core.modified = datetime(2026, 4, 15, 9, 0, 0)
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(9.5)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = header.add_run("HR OPS CONTROL · 15 APRIL 2026")
    hrun.font.name = "Aptos"
    hrun.font.size = Pt(8)
    hrun.font.color.rgb = RGBColor(89, 89, 89)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = footer.add_run("Frozen 2026-04-15 · Draft routing only · No message or access side effect")
    frun.font.name = "Aptos"
    frun.font.size = Pt(8)

    add_title(document, "Compliance Training Enforcement Brief", "Policy-controlled audit package · frozen 15 April 2026")
    add_heading(document, "Executive control result")
    add_body(
        document,
        "Eight source rows were reconciled against Compliance Handbook 2026 Section 6 and the frozen date. The controlled result is 2 WARNING, 1 SUSPENSION_REQUEST, 3 COMMITTEE_ESCALATION, 1 EXEMPT and 1 NO_ACTION. Six cases are routed; two are excluded by an active exemption or completed status."
    )
    metrics = document.add_table(rows=2, cols=6)
    labels = ["Cases", "Warning", "Suspend request", "Committee", "Exempt", "No action"]
    values = ["8", "2", "1", "3", "1", "1"]
    for index, label in enumerate(labels):
        metrics.cell(0, index).text = label
        metrics.cell(1, index).text = values[index]
    format_doc_table(metrics, font_size=8.5)
    add_heading(document, "Decision order")
    add_body(
        document,
        "Apply source completion status first, then an active exemption, then overdue bands: 1–14 days warning to employee with manager CC; 15–30 days IT suspension request plus manager notice; over 30 days committee-only escalation with no direct action. The committee must not receive 1–30 day cases."
    )
    add_heading(document, "Evidence chain")
    add_body(
        document,
        "Training_Compliance_Tracker.xlsx supplies eight source rows. Compliance_Handbook.mbox / msg_handbook_5106 supplies Section 6. Workflow_Context.json freezes 2026-04-15T09:00:00Z. Compliance_Channel.json identifies C_COMP_5106 / compliance."
    )
    add_heading(document, "Side-effect boundary")
    add_body(
        document,
        "Five email artifacts and one channel artifact are drafts only. Access_Suspension_Control.json records one request for Samantha Green but executed=false and execution_record_count=0. No email was sent, no Slack message was posted and no system access was changed."
    )

    document.add_page_break()
    add_title(document, "Case Register", "Source status -> date/exemption precedence -> controlled route")
    table = document.add_table(rows=1, cols=7)
    headers = ["Employee", "Training", "Due", "Days", "Override", "Action", "Route"]
    for index, value in enumerate(headers):
        table.cell(0, index).text = value
    for item in decisions:
        cells = table.add_row().cells
        override = "Complete" if item.completion_override else ("Active exemption" if item.active_exemption else "None")
        route = {
            "WARNING": "Employee + manager CC",
            "SUSPENSION_REQUEST": "IT Security + manager",
            "COMMITTEE_ESCALATION": "Committee only",
            "EXEMPT": "No notification",
            "NO_ACTION": "No notification",
        }[item.action]
        values = [
            item.source.employee, item.source.training, item.source.due_date, str(item.days_overdue),
            override, item.action, route,
        ]
        for index, value in enumerate(values):
            cells[index].text = value
    format_doc_table(table, widths=[3.8, 4.0, 2.5, 1.6, 3.5, 4.4, 7.0], font_size=7.7)
    add_heading(document, "Identity and boundary checks")
    add_body(
        document,
        "Patricia Chen is the 14-day warning case; Patricia Chang is the 36-day committee case. Roberto Alvarez remains exempt through April 30 despite 45 overdue days. Wendy Torres remains NO_ACTION because the source status is Complete. Theodore Kim and Victor Pham are committee-only and must not receive direct warnings or IT requests."
    )
    add_heading(document, "Case-specific routing facts")
    facts = document.add_table(rows=1, cols=3)
    for index, value in enumerate(["Cases", "Required draft", "Forbidden route"]):
        facts.cell(0, index).text = value
    rows = [
        ("Patricia Chen; Ursula Martin", "Individual warning; manager CC", "Compliance committee"),
        ("Samantha Green", "IT request + Diana Torres notice", "Committee; claim of execution"),
        ("Patricia Chang; Theodore Kim; Victor Pham", "One committee-only escalation", "Employee or IT direct action"),
        ("Roberto Alvarez; Wendy Torres", "No draft", "Any recipient"),
    ]
    for row in rows:
        cells = facts.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    format_doc_table(facts, widths=[8.0, 8.4, 8.4], font_size=8.0)

    document.add_page_break()
    add_title(document, "Routing & Control Register", "Five drafts, aggregate channel summary and zero execution")
    add_heading(document, "Offline routing drafts")
    routing = document.add_table(rows=1, cols=5)
    for index, value in enumerate(["Draft", "To", "CC", "Cases", "Status"]):
        routing.cell(0, index).text = value
    routing_rows = [
        ("Warning 1", "patricia.chen@company.example.com", "vikram.shah@company.example.com", "Patricia Chen · 14 days", "DRAFT_NOT_SENT"),
        ("Warning 2", "ursula.martin@company.example.com", "diana.torres@company.example.com", "Ursula Martin · 5 days", "DRAFT_NOT_SENT"),
        ("IT request", "it-security@company.example.com", "—", "Samantha Green · 21 days", "DRAFT_NOT_SENT"),
        ("Manager", "diana.torres@company.example.com", "—", "Samantha Green · 21 days", "DRAFT_NOT_SENT"),
        ("Committee", "compliance-committee@company.example.com", "—", "Patricia Chang 36; Theodore Kim 64; Victor Pham 121", "DRAFT_NOT_SENT"),
    ]
    for row in routing_rows:
        cells = routing.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    format_doc_table(routing, widths=[2.8, 6.3, 6.0, 7.2, 3.6], font_size=7.6)
    add_heading(document, "Section 6 policy register")
    policy_table = document.add_table(rows=1, cols=3)
    for index, value in enumerate(["Rule", "Action", "Boundary"]):
        policy_table.cell(0, index).text = value
    policy_rows = [
        ("1–14 days", "Employee warning + manager CC", "No committee"),
        ("15–30 days", "IT suspension request + manager notice", "No committee; request is not execution"),
        (">30 days", "Compliance committee escalation", "No direct action"),
        ("Active exemption", "Exclude from enforcement", "Overrides overdue"),
    ]
    for row in policy_rows:
        cells = policy_table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    format_doc_table(policy_table, widths=[5.0, 10.5, 9.3], font_size=8.0)
    add_heading(document, "Privacy and final attestation")
    add_body(
        document,
        "The C_COMP_5106 / compliance draft reports only aggregate 8/2/1/3/1/1 counts and contains no name, email, training, due date or case-level overdue days. Status is DRAFT_NOT_POSTED and posted=false. All five emails remain DRAFT_NOT_SENT. Access suspension executed: false; execution records: 0."
    )

    document.save(path)
    normalize_ooxml(path)


def draft_message(to: str, subject: str, body: str, message_id: str, cc: str = "") -> EmailMessage:
    message = EmailMessage(policy=SMTP)
    message["From"] = "hr-ops@company.example.com"
    message["To"] = to
    if cc:
        message["Cc"] = cc
    message["Subject"] = subject
    message["Date"] = "Wed, 15 Apr 2026 09:00:00 +0000"
    message["Message-ID"] = f"<{message_id}@evocowork.invalid>"
    message["X-EvoCowork-Status"] = "DRAFT_NOT_SENT"
    message["X-EvoCowork-Source"] = "AutomationBench-5106"
    message.set_content(body.rstrip() + "\n", charset="utf-8")
    return message


def write_mbox(path: Path, decisions: list[Decision]) -> None:
    by_name = {item.source.employee: item for item in decisions}
    messages: list[EmailMessage] = []
    for sequence, name in enumerate(("Patricia Chen", "Ursula Martin"), start=1):
        item = by_name[name]
        body = (
            f"DRAFT_NOT_SENT\n\n{name},\n\n"
            f"The {item.source.training} training was due {item.source.due_date} and is {item.days_overdue} days overdue as of 2026-04-15. "
            "Compliance Handbook 2026 Section 6 routes 1–14 day cases to an employee warning with manager CC.\n\n"
            "This is an offline warning draft. It has not been sent, and no access action has been taken."
        )
        messages.append(draft_message(
            item.source.employee_email, f"DRAFT — overdue compliance training — {name}", body,
            f"automationbench-5106-warning-{sequence}", cc=item.source.manager_email,
        ))

    samantha = by_name["Samantha Green"]
    it_body = (
        "DRAFT_NOT_SENT\n\nIT Security,\n\n"
        f"Please review an access-suspension request for Samantha Green ({samantha.source.employee_email}). "
        f"Her {samantha.source.training} training was due {samantha.source.due_date} and is {samantha.days_overdue} days overdue as of 2026-04-15. "
        "Compliance Handbook 2026 Section 6 routes 15–30 day cases to IT.\n\n"
        "This is an offline request draft only. It has not been sent and access has not been suspended."
    )
    messages.append(draft_message(
        "it-security@company.example.com", "DRAFT — access suspension request — Samantha Green",
        it_body, "automationbench-5106-it-samantha",
    ))
    manager_body = (
        "DRAFT_NOT_SENT\n\nDiana Torres,\n\n"
        f"Samantha Green's {samantha.source.training} training was due {samantha.source.due_date} and is {samantha.days_overdue} days overdue as of 2026-04-15. "
        "Under Compliance Handbook 2026 Section 6, an IT access-suspension request draft has been prepared.\n\n"
        "This manager notice is an offline draft. It has not been sent and access has not been suspended."
    )
    messages.append(draft_message(
        samantha.source.manager_email, "DRAFT — manager notice — Samantha Green compliance training",
        manager_body, "automationbench-5106-manager-samantha",
    ))

    committee_names = ("Patricia Chang", "Theodore Kim", "Victor Pham")
    lines = []
    for name in committee_names:
        item = by_name[name]
        lines.append(f"- {name} | {item.source.training} | due {item.source.due_date} | {item.days_overdue} days overdue")
    committee_body = (
        "DRAFT_NOT_SENT\n\nCompliance Committee,\n\n"
        "The following cases exceed 30 days and require committee-only escalation under Compliance Handbook 2026 Section 6:\n"
        + "\n".join(lines)
        + "\n\nNo 1–30 day case is included. This is one consolidated offline draft. It has not been sent; no direct warning, IT request or access action was taken for these cases."
    )
    messages.append(draft_message(
        "compliance-committee@company.example.com", "DRAFT — over-30-day compliance escalations",
        committee_body, "automationbench-5106-committee",
    ))

    chunks: list[bytes] = []
    for message in messages:
        chunks.append(b"From hr-ops@company.example.com Wed Apr 15 09:00:00 2026\n")
        buffer = BytesIO()
        BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
        chunks.append(buffer.getvalue().replace(b"\r\n", b"\n"))
        chunks.append(b"\n")
    path.write_bytes(b"".join(chunks))


def write_slack(path: Path) -> None:
    payload = {
        "source_channel": {"id": "C_COMP_5106", "name": "compliance", "is_private": False},
        "status": "DRAFT_NOT_POSTED",
        "posted": False,
        "summary": {
            "total_source_cases": 8,
            "warning": 2,
            "suspension_request": 1,
            "committee_escalation": 3,
            "exempt": 1,
            "no_action": 1,
        },
        "draft_message": {
            "text": "Training compliance control: 8 source cases reviewed; 2 WARNING, 1 SUSPENSION_REQUEST, 3 COMMITTEE_ESCALATION, 1 EXEMPT and 1 NO_ACTION. Five routing drafts prepared; no message or access action executed."
        },
        "privacy_control": {
            "aggregate_only": True,
            "employee_names_included": False,
            "employee_emails_included": False,
            "training_titles_included": False,
            "due_dates_included": False,
            "case_level_days_included": False,
        },
        "source_trace": [
            "Training_Enforcement_Audit.xlsx:Control_Summary",
            "Compliance_Channel.json:channels[0]",
        ],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_execution_control(path: Path) -> None:
    payload = {
        "decision": "DRAFT_REQUEST_ONLY",
        "status": "DRAFT_REQUEST_ONLY",
        "executed": False,
        "execution_record_count": 0,
        "transaction_ids": [],
        "request_count": 1,
        "requests": [
            {
                "employee": "Samantha Green",
                "training": "Information Security",
                "due_date": "2026-03-25",
                "days_overdue": 21,
                "recipient": "it-security@company.example.com",
                "source_row": 5,
                "status": "DRAFT_NOT_SENT",
            }
        ],
        "control_reason": "HR Ops prepares the Section 6 request; the benchmark does not execute an access change.",
        "source_policy": {
            "file": "Compliance_Handbook.mbox",
            "message_id": "msg_handbook_5106",
            "section": "Section 6: Enforcement Actions",
            "rule": "15 to 30 days overdue",
        },
        "source_trace": [
            "Training_Compliance_Tracker.xlsx:Training Status row 5",
            "Compliance_Handbook.mbox:msg_handbook_5106:Section 6",
            "Training_Enforcement_Audit.xlsx:Notification_Queue D003",
        ],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=Path("/app/input_files"))
    parser.add_argument("--output-dir", type=Path, default=Path("/app/output"))
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    if any(args.output_dir.iterdir()):
        raise ValueError("输出目录必须为空，避免遗留文件污染交付")

    cases, handbook, channel, frozen_time = read_inputs(args.input_dir)
    decisions = derive_decisions(cases)
    write_workbook(args.output_dir / XLSX_NAME, cases, decisions, handbook, channel, frozen_time)
    write_brief(args.output_dir / DOCX_NAME, decisions)
    write_mbox(args.output_dir / MBOX_NAME, decisions)
    write_slack(args.output_dir / SLACK_NAME)
    write_execution_control(args.output_dir / CONTROL_NAME)


if __name__ == "__main__":
    main()
