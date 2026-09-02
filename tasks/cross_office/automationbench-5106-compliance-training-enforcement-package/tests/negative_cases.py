#!/usr/bin/env python3
"""对 AutomationBench 5106 Golden 逐项注入本题专属错误并验证 Judge 拒绝。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Callable

from docx import Document
from lxml import etree
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
XLSX = "Training_Enforcement_Audit.xlsx"
DOCX = "Compliance_Enforcement_Brief.docx"
MBOX = "Enforcement_Notification_Drafts.mbox"
SLACK = "Compliance_Channel_Draft.json"
CONTROL = "Access_Suspension_Control.json"
Mutator = Callable[[Path], None]


def load_judge():
    spec = importlib.util.spec_from_file_location("compliance_5106_judge", ROOT / "tests" / "judge.py")
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入本题 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def rewrite_zip(path: Path, transform: Callable[[str, bytes], bytes | None]) -> None:
    temporary = path.with_suffix(path.suffix + ".negative")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = transform(info.filename, source.read(info.filename))
            if data is not None:
                target.writestr(info, data)
    temporary.replace(path)


def replace_ooxml_text(path: Path, old: str, new: str) -> None:
    old_bytes, new_bytes = old.encode("utf-8"), new.encode("utf-8")
    changed = False

    def transform(name: str, data: bytes) -> bytes:
        nonlocal changed
        if name.endswith(".xml") and old_bytes in data:
            changed = True
            return data.replace(old_bytes, new_bytes)
        return data

    rewrite_zip(path, transform)
    if not changed:
        raise ValueError(f"OOXML 中找不到 {old!r}")


def sheet_xml_path(archive: zipfile.ZipFile, sheet_name: str) -> str:
    workbook = etree.fromstring(archive.read("xl/workbook.xml"))
    relationships = etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    ns = {
        "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    rel_id = workbook.xpath(f'//m:sheet[@name="{sheet_name}"]/@r:id', namespaces=ns)[0]
    target = relationships.xpath(f'//pr:Relationship[@Id="{rel_id}"]/@Target', namespaces=ns)[0]
    if target.startswith("/"):
        return target.lstrip("/")
    return "xl/" + target.lstrip("/")


def edit_cell(path: Path, sheet_name: str, coordinate: str, *, value: str | int | None = None, remove_formula: bool = False) -> None:
    with zipfile.ZipFile(path, "r") as archive:
        target_sheet = sheet_xml_path(archive, sheet_name)
    changed = False
    ns_uri = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

    def transform(name: str, data: bytes) -> bytes:
        nonlocal changed
        if name != target_sheet:
            return data
        root = etree.fromstring(data)
        cells = root.xpath(f'//m:c[@r="{coordinate}"]', namespaces={"m": ns_uri})
        if not cells:
            raise ValueError(f"{sheet_name}!{coordinate} 不存在")
        cell = cells[0]
        if remove_formula:
            formula = cell.find(f"{{{ns_uri}}}f")
            if formula is None:
                raise ValueError(f"{sheet_name}!{coordinate} 没有公式")
            cell.remove(formula)
        if value is not None:
            node = cell.find(f"{{{ns_uri}}}v")
            if node is None:
                node = etree.SubElement(cell, f"{{{ns_uri}}}v")
            node.text = str(value)
            if isinstance(value, int):
                cell.attrib.pop("t", None)
            else:
                cell.set("t", "str")
        changed = True
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    rewrite_zip(path, transform)
    if not changed:
        raise ValueError("目标单元格未修改")


def mutate_json(path: Path, mutator: Callable[[dict], None]) -> None:
    payload = json.loads(path.read_text(encoding="utf-8"))
    mutator(payload)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def replace_bytes(path: Path, old: bytes, new: bytes) -> None:
    data = path.read_bytes()
    if old not in data:
        raise ValueError(f"MBOX 中找不到 {old!r}")
    path.write_bytes(data.replace(old, new, 1))


def drop_message(path: Path, token: bytes) -> None:
    chunks = [chunk for chunk in re.split(br"(?m)(?=^From hr-ops@company\.example\.com )", path.read_bytes()) if chunk]
    kept = [chunk for chunk in chunks if token not in chunk]
    if len(kept) != len(chunks) - 1:
        raise ValueError(f"无法唯一删除包含 {token!r} 的邮件")
    path.write_bytes(b"".join(kept))


def replace_docx_text(path: Path, old: str, new: str, *, row_token: str | None = None) -> None:
    document = Document(path)
    changed = False
    paragraphs = list(document.paragraphs) if row_token is None else []
    for table in document.tables:
        for row in table.rows:
            joined = " | ".join(cell.text for cell in row.cells)
            if row_token is not None and row_token not in joined:
                continue
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        if old in paragraph.text:
            paragraph.text = paragraph.text.replace(old, new, 1)
            changed = True
            break
    if not changed:
        raise ValueError(f"Word 中找不到 {old!r}")
    document.save(path)


def openpyxl_mutation(path: Path, mutator: Callable) -> None:
    workbook = load_workbook(path)
    mutator(workbook)
    workbook.save(path)


def build_cases() -> list[tuple[str, Mutator]]:
    cases: list[tuple[str, Mutator]] = []
    cases.extend([
        ("missing-docx", lambda root: (root / DOCX).unlink()),
        ("extra-debug-file", lambda root: (root / "debug.txt").write_text("debug\n", encoding="utf-8")),
        ("corrupt-xlsx", lambda root: (root / XLSX).write_bytes(b"not-an-xlsx")),
        ("source-name-changed", lambda root: replace_ooxml_text(root / XLSX, "Patricia Chen", "Patricia Chan")),
        ("source-due-date-changed", lambda root: replace_ooxml_text(root / XLSX, "2026-04-01", "2026-04-02")),
        ("source-roberto-note-changed", lambda root: replace_ooxml_text(root / XLSX, "Medical leave exemption through April 30", "Medical leave ended March 30")),
        ("source-wendy-status-changed", lambda root: replace_ooxml_text(root / XLSX, "Complete", "Overdue")),
        ("policy-warning-band-changed", lambda root: replace_ooxml_text(root / XLSX, "1 to 14 days overdue", "1 to 10 days overdue")),
        ("policy-committee-restriction-changed", lambda root: replace_ooxml_text(root / XLSX, "Do NOT take direct action", "Take direct action")),
        ("evidence-time-changed", lambda root: replace_ooxml_text(root / XLSX, "2026-04-15T09:00:00Z", "2026-04-16T09:00:00Z")),
        ("evidence-channel-changed", lambda root: replace_ooxml_text(root / XLSX, "C_COMP_5106", "C_COMP_WRONG")),
        ("formula-days-staticized", lambda root: edit_cell(root / XLSX, "Enforcement_Decisions", "J2", remove_formula=True)),
        ("formula-action-staticized", lambda root: edit_cell(root / XLSX, "Enforcement_Decisions", "M5", remove_formula=True)),
        ("formula-exemption-staticized", lambda root: edit_cell(root / XLSX, "Enforcement_Decisions", "L4", remove_formula=True)),
        ("chen-days-wrong", lambda root: edit_cell(root / XLSX, "Enforcement_Decisions", "J2", value=15)),
        ("chang-action-warning", lambda root: edit_cell(root / XLSX, "Enforcement_Decisions", "M3", value="WARNING")),
        ("roberto-action-escalation", lambda root: edit_cell(root / XLSX, "Enforcement_Decisions", "M4", value="COMMITTEE_ESCALATION")),
        ("samantha-action-committee", lambda root: edit_cell(root / XLSX, "Enforcement_Decisions", "M5", value="COMMITTEE_ESCALATION")),
        ("theodore-action-suspension", lambda root: edit_cell(root / XLSX, "Enforcement_Decisions", "M6", value="SUSPENSION_REQUEST")),
        ("ursula-days-wrong", lambda root: edit_cell(root / XLSX, "Enforcement_Decisions", "J7", value=15)),
        ("victor-action-suspension", lambda root: edit_cell(root / XLSX, "Enforcement_Decisions", "M8", value="SUSPENSION_REQUEST")),
        ("wendy-action-warning", lambda root: edit_cell(root / XLSX, "Enforcement_Decisions", "M9", value="WARNING")),
        ("summary-warning-count-wrong", lambda root: edit_cell(root / XLSX, "Control_Summary", "B3", value=3)),
        ("summary-draft-count-wrong", lambda root: edit_cell(root / XLSX, "Control_Summary", "B9", value=4)),
        ("queue-committee-case-missing", lambda root: edit_cell(root / XLSX, "Notification_Queue", "E6", value="Patricia Chang, Theodore Kim")),
        ("queue-it-wrong-case", lambda root: edit_cell(root / XLSX, "Notification_Queue", "E4", value="Theodore Kim")),
        ("native-chart-removed", lambda root: openpyxl_mutation(root / XLSX, lambda wb: setattr(wb["Control_Summary"], "_charts", []))),
        ("native-table-removed", lambda root: openpyxl_mutation(root / XLSX, lambda wb: wb["Source_Training"].tables.clear())),
        ("conditional-formatting-removed", lambda root: openpyxl_mutation(root / XLSX, lambda wb: setattr(wb["Enforcement_Decisions"], "conditional_formatting", type(wb["Enforcement_Decisions"].conditional_formatting)()))),
    ])

    cases.extend([
        ("mbox-drop-committee", lambda root: drop_message(root / MBOX, b"over-30-day compliance escalations")),
        ("mbox-drop-chen", lambda root: drop_message(root / MBOX, b"Patricia Chen")),
        ("mbox-chen-wrong-cc", lambda root: replace_bytes(root / MBOX, b"Cc: vikram.shah@company.example.com", b"Cc: diana.torres@company.example.com")),
        ("mbox-chen-committee-cc", lambda root: replace_bytes(root / MBOX, b"Cc: vikram.shah@company.example.com", b"Cc: compliance-committee@company.example.com")),
        ("mbox-ursula-wrong-recipient", lambda root: replace_bytes(root / MBOX, b"To: ursula.martin@company.example.com", b"To: roberto.alvarez@company.example.com")),
        ("mbox-it-direct-to-samantha", lambda root: replace_bytes(root / MBOX, b"To: it-security@company.example.com", b"To: samantha.green@company.example.com")),
        ("mbox-it-claim-executed", lambda root: replace_bytes(root / MBOX, b"access has no=\nt been suspended", b"access has be=\nen suspended")),
        ("mbox-manager-wrong-recipient", lambda root: replace_bytes(root / MBOX, b"To: diana.torres@company.example.com", b"To: vikram.shah@company.example.com")),
        ("mbox-committee-drop-patricia", lambda root: replace_bytes(root / MBOX, b"- Patricia Chang | Anti-Harassment | due 2026-03-10 | 36 days overdue\n", b"")),
        ("mbox-committee-drop-theodore", lambda root: replace_bytes(root / MBOX, b"- Theodore Kim | Anti-Harassment | due 2026-02-10 | 64 days overdue\n", b"")),
        ("mbox-committee-drop-victor", lambda root: replace_bytes(root / MBOX, b"- Victor Pham | Information Security | due 2025-12-15 | 121 days overdue\n", b"")),
        ("mbox-committee-add-chen", lambda root: replace_bytes(root / MBOX, b"The following cases exceed 30 days", b"Patricia Chen; the following cases exceed 30 days")),
        ("mbox-false-sent", lambda root: replace_bytes(root / MBOX, b"X-EvoCowork-Status: DRAFT_NOT_SENT", b"X-EvoCowork-Status: SENT")),
        ("mbox-roberto-recipient", lambda root: replace_bytes(root / MBOX, b"To: patricia.chen@company.example.com", b"To: roberto.alvarez@company.example.com")),
    ])

    cases.extend([
        ("slack-false-posted", lambda root: mutate_json(root / SLACK, lambda p: p.update({"posted": True, "status": "POSTED"}))),
        ("slack-wrong-channel", lambda root: mutate_json(root / SLACK, lambda p: p["source_channel"].update({"id": "C_WRONG"}))),
        ("slack-warning-count-wrong", lambda root: mutate_json(root / SLACK, lambda p: p["summary"].update({"warning": 3}))),
        ("slack-committee-count-wrong", lambda root: mutate_json(root / SLACK, lambda p: p["summary"].update({"committee_escalation": 2}))),
        ("slack-leak-name", lambda root: mutate_json(root / SLACK, lambda p: p["draft_message"].update({"text": p["draft_message"]["text"] + " Patricia Chen"}))),
        ("slack-leak-email", lambda root: mutate_json(root / SLACK, lambda p: p["draft_message"].update({"text": p["draft_message"]["text"] + " samantha.green@company.example.com"}))),
        ("slack-leak-training", lambda root: mutate_json(root / SLACK, lambda p: p["draft_message"].update({"text": p["draft_message"]["text"] + " Information Security"}))),
        ("slack-leak-due-date", lambda root: mutate_json(root / SLACK, lambda p: p["draft_message"].update({"text": p["draft_message"]["text"] + " 2026-03-25"}))),
        ("slack-privacy-name-flag", lambda root: mutate_json(root / SLACK, lambda p: p["privacy_control"].update({"employee_names_included": True}))),
        ("slack-source-trace-missing", lambda root: mutate_json(root / SLACK, lambda p: p.update({"source_trace": []}))),
    ])

    cases.extend([
        ("control-false-executed", lambda root: mutate_json(root / CONTROL, lambda p: p.update({"executed": True, "execution_record_count": 1}))),
        ("control-fake-transaction", lambda root: mutate_json(root / CONTROL, lambda p: p.update({"transaction_ids": ["TX-FAKE"]}))),
        ("control-wrong-employee", lambda root: mutate_json(root / CONTROL, lambda p: p["requests"][0].update({"employee": "Theodore Kim"}))),
        ("control-wrong-training", lambda root: mutate_json(root / CONTROL, lambda p: p["requests"][0].update({"training": "Anti-Harassment"}))),
        ("control-wrong-days", lambda root: mutate_json(root / CONTROL, lambda p: p["requests"][0].update({"days_overdue": 31}))),
        ("control-wrong-source-row", lambda root: mutate_json(root / CONTROL, lambda p: p["requests"][0].update({"source_row": 6}))),
        ("control-wrong-status", lambda root: mutate_json(root / CONTROL, lambda p: p.update({"status": "EXECUTED"}))),
        ("control-wrong-policy", lambda root: mutate_json(root / CONTROL, lambda p: p["source_policy"].update({"rule": "Over 30 days overdue"}))),
        ("control-extra-request", lambda root: mutate_json(root / CONTROL, lambda p: (p.update({"request_count": 2}), p["requests"].append(dict(p["requests"][0]))))),
        ("control-source-trace-missing", lambda root: mutate_json(root / CONTROL, lambda p: p.update({"source_trace": []}))),
    ])

    cases.extend([
        ("word-total-wrong", lambda root: replace_docx_text(root / DOCX, "Eight source rows", "Seven source rows")),
        ("word-chen-action-wrong", lambda root: replace_docx_text(root / DOCX, "WARNING", "COMMITTEE_ESCALATION", row_token="Patricia Chen")),
        ("word-chang-days-wrong", lambda root: replace_docx_text(root / DOCX, "36", "14", row_token="Patricia Chang")),
        ("word-roberto-action-wrong", lambda root: replace_docx_text(root / DOCX, "EXEMPT", "COMMITTEE_ESCALATION", row_token="Roberto Alvarez")),
        ("word-samantha-action-wrong", lambda root: replace_docx_text(root / DOCX, "SUSPENSION_REQUEST", "WARNING", row_token="Samantha Green")),
        ("word-theodore-days-wrong", lambda root: replace_docx_text(root / DOCX, "64", "24", row_token="Theodore Kim")),
        ("word-ursula-action-wrong", lambda root: replace_docx_text(root / DOCX, "WARNING", "SUSPENSION_REQUEST", row_token="Ursula Martin")),
        ("word-victor-action-wrong", lambda root: replace_docx_text(root / DOCX, "COMMITTEE_ESCALATION", "SUSPENSION_REQUEST", row_token="Victor Pham")),
        ("word-wendy-action-wrong", lambda root: replace_docx_text(root / DOCX, "NO_ACTION", "WARNING", row_token="Wendy Torres")),
        ("word-false-access-execution", lambda root: replace_docx_text(root / DOCX, "executed=false", "executed=true")),
        ("word-policy-removed", lambda root: replace_docx_text(root / DOCX, "Active exemption", "No exemption", row_token="Exclude from enforcement")),
    ])
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    judge = load_judge()
    # Golden 的真实 LibreOffice 渲染已单独验收；负例只测试语义与结构拒错，
    # 避免每个 JSON/MBOX 变体重复启动 Office 进程。
    judge.render_office = lambda path: (6 if path.suffix.casefold() == ".xlsx" else 3, "")
    results: list[dict] = []
    for name, mutator in build_cases():
        with tempfile.TemporaryDirectory(prefix=f"negative-5106-{name}-") as temporary:
            output = Path(temporary) / "output"
            shutil.copytree(GOLDEN, output)
            mutator(output)
            result = judge.evaluate(output)
            results.append({
                "case": name,
                "reward": result.get("reward", 0.0),
                "passed": result.get("passed", False),
                "hurdle_failures": result.get("hurdle_failures", []),
                "penalties": result.get("penalties", []),
            })
    report = {
        "total": len(results),
        "rejected": sum(not item["passed"] for item in results),
        "maximum_reward": max(item["reward"] for item in results),
        "cases": results,
    }
    if args.report:
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if report["rejected"] != report["total"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
