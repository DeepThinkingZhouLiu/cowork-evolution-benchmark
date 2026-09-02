#!/usr/bin/env python3
"""只把 AutomationBench 5107 的公开初始状态静态化为本题输入。

本脚本只接受 ``hr.equipment_provisioning``，逐项断言八名新入职人员、
完整配置政策、冻结时间、IT 频道和九条上游断言。它不是通用制题器，
不会创造人员、设备型号、地点、雇佣类型或采购事实。
"""

from __future__ import annotations

import importlib.util
import json
import shutil
import sys
import types
import zipfile
from datetime import datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
UPSTREAM = Path(__file__).with_name("upstream-hr-tasks.py")
SNAPSHOT = Path(__file__).with_name("upstream-task-5107.json")
DATA_DIR = ROOT / "data" / "input_files"
ENV_DIR = ROOT / "environment" / "input_files"
FIXED_ZIP_TIME = (2026, 4, 15, 9, 0, 0)


def load_task() -> dict:
    datasets = types.ModuleType("datasets")
    datasets.Dataset = object
    sys.modules["datasets"] = datasets
    for name in ("automationbench", "automationbench.domains", "automationbench.domains.hr"):
        sys.modules[name] = types.ModuleType(name)
    noise = types.ModuleType("automationbench.domains.hr._noise")
    noise.apply_noise = lambda value: value
    sys.modules[noise.__name__] = noise
    spec = importlib.util.spec_from_file_location("automationbench_hr_tasks_frozen_5107", UPSTREAM)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入冻结的 AutomationBench HR 源文件")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    task = module.get_task_5107()
    if task.get("example_id") != 5107 or task.get("task") != "hr.equipment_provisioning":
        raise ValueError("冻结函数不再是 AutomationBench 5107")
    return task


def assert_shape(task: dict) -> None:
    state = task["info"]["initial_state"]
    assert state["meta"] == {"schema_version": "0.1.0", "current_time": "2026-04-15T09:00:00Z"}
    messages = state["gmail"]["messages"]
    assert [item["id"] for item in messages] == ["msg_provision_policy_5107"]
    body = messages[0]["body_plain"]
    assert all(token in body for token in (
        "Contractors and consultants must provide their own equipment",
        "Remote employees get an additional monitor",
        "MacBook Pro 16-inch M4", "MacBook Air 15-inch M4",
    ))
    spreadsheets = state["google_sheets"]["spreadsheets"]
    assert len(spreadsheets) == 1 and spreadsheets[0]["id"] == "ss_newhires_5107"
    worksheet = spreadsheets[0]["worksheets"][0]
    assert worksheet["title"] == "April Starters"
    assert [item["row_id"] for item in worksheet["rows"]] == list(range(2, 10))
    assert [item["cells"]["Name"] for item in worksheet["rows"]] == [
        "Ada Petrova", "Brian Osei", "Carmen Delgado", "David Nguyen",
        "Elena Rossi", "Femi Adeyemi", "Greta Holm", "Hugo Fernandez",
    ]
    assert worksheet["rows"][4]["cells"]["Notes"] == "Equipment already shipped from previous office transfer"
    assert state["slack"]["channels"] == [{"id": "C_IT_5107", "name": "it-general", "is_private": False}]
    assert state["slack"]["messages"] == []
    assert len(task["info"]["assertions"]) == 9


def normalize_xlsx(path: Path) -> None:
    temporary = path.with_suffix(".normalized.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    temporary.replace(path)


def write_tracker(state: dict, target: Path) -> None:
    spreadsheet = state["google_sheets"]["spreadsheets"][0]
    worksheet = spreadsheet["worksheets"][0]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = worksheet["title"]
    workbook.creator = "AutomationBench 5107 public fixture"
    workbook.title = spreadsheet["title"]
    workbook.subject = "Static Office representation of April starters"
    workbook.description = "No person, date, location, employment type or note was added beyond AutomationBench 5107."
    headers = ["Source Row", "Name", "Employee ID", "Department", "Start Date", "Location", "Employment Type", "Notes"]
    sheet.append(headers)
    for item in worksheet["rows"]:
        cells = item["cells"]
        sheet.append([item["row_id"]] + [cells.get(header, "") for header in headers[1:]])
    table = Table(displayName="AprilStartersSourceInput", ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    widths = {"A": 12, "B": 23, "C": 17, "D": 20, "E": 17, "F": 25, "G": 25, "H": 58}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = sheet.dimensions
    workbook.save(target)
    normalize_xlsx(target)


def message_bytes(record: dict) -> bytes:
    message = EmailMessage(policy=SMTP)
    message["From"] = record["from_"]
    message["To"] = ", ".join(record["to"])
    message["Subject"] = record["subject"]
    parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
    message["Date"] = parsed.strftime("%a, %d %b %Y %H:%M:%S +0000")
    message["Message-ID"] = f"<{record['id']}@automationbench.invalid>"
    message["X-AutomationBench-Message-ID"] = record["id"]
    message["X-AutomationBench-Thread-ID"] = record["thread_id"]
    message["X-AutomationBench-Labels"] = ",".join(record.get("label_ids", []))
    message.set_content(record["body_plain"], charset="utf-8")
    buffer = BytesIO()
    BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
    return buffer.getvalue()


def write_mbox(messages: list[dict], target: Path) -> None:
    chunks: list[bytes] = []
    for record in messages:
        parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
        chunks.append(f"From {record['from_']} {parsed.strftime('%a %b %d %H:%M:%S %Y')}\n".encode("ascii"))
        chunks.append(message_bytes(record).replace(b"\r\n", b"\n"))
        chunks.append(b"\n")
    target.write_bytes(b"".join(chunks))


def main() -> None:
    task = load_task()
    assert_shape(task)
    state = task["info"]["initial_state"]
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    ENV_DIR.mkdir(parents=True, exist_ok=True)
    SNAPSHOT.write_text(json.dumps(task, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_tracker(state, DATA_DIR / "New_Hire_Onboarding_Tracker.xlsx")
    write_mbox(state["gmail"]["messages"], DATA_DIR / "Provisioning_Policy_Mailbox.mbox")
    (DATA_DIR / "IT_Channel.json").write_text(json.dumps(state["slack"], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (DATA_DIR / "Workflow_Context.json").write_text(json.dumps({"meta": state["meta"]}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    for source in sorted(DATA_DIR.iterdir()):
        if source.is_file():
            shutil.copy2(source, ENV_DIR / source.name)


if __name__ == "__main__":
    main()
