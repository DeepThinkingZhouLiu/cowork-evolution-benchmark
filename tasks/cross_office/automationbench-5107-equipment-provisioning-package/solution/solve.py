#!/usr/bin/env python3
"""AutomationBench 5107 入职设备配置包的逐题参考解法。"""

from __future__ import annotations

import argparse
import hashlib
import json
import mailbox
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


INPUT_HASHES = {
    "New_Hire_Onboarding_Tracker.xlsx": "aa8ac492cf0995a480de8ff5a9435c19b601a457ad4d358f029c9b0525e0d20d",
    "Provisioning_Policy_Mailbox.mbox": "a9b0849b12b6b4642a7ffea9227ba39fe2460ea43435bd33a4f07fb7137e6124",
    "IT_Channel.json": "c98d2dbebd2d0831d2b82e5309bd71fbfd7cb3bc6e718fcc502d5d27432fab32",
    "Workflow_Context.json": "76c4a0657bc27f2fd830b10613ec7ceae3adcdc5798aeee0f822d27bf69bbc99",
}
XLSX_NAME = "Equipment_Provisioning_Audit.xlsx"
DOCX_NAME = "Provisioning_Control_Brief.docx"
JIRA_NAME = "Provisioning_Jira_Drafts.json"
MBOX_NAME = "IT_Equipment_List_Draft.mbox"
SLACK_NAME = "IT_Channel_Draft.json"
CONTROL_NAME = "Purchase_Execution_Control.json"
FIXED_ISO_TIME = "2026-04-15T09:00:00Z"
FIXED_ZIP_TIME = (2026, 4, 15, 9, 0, 0)

NAVY = "17365D"
GREEN = "E2F0D9"
AMBER = "FFF2CC"
RED = "F4CCCC"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="B7C9D6")


@dataclass(frozen=True)
class Starter:
    source_row: int
    name: str
    employee_id: str
    department: str
    start_date: str
    location: str
    employment_type: str
    notes: str


@dataclass(frozen=True)
class Decision:
    source: Starter
    contractor: bool
    already_provisioned: bool
    status: str
    laptop: str
    monitor: str
    keyboard: str
    mouse: str
    headset: str
    jira_plan: str
    email_list: str


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def decode_message(message) -> str:
    if message.is_multipart():
        parts: list[str] = []
        for part in message.walk():
            if part.get_content_type() == "text/plain":
                payload = part.get_payload(decode=True) or b""
                parts.append(payload.decode(part.get_content_charset() or "utf-8", errors="replace"))
        return "\n".join(parts).strip()
    payload = message.get_payload(decode=True) or b""
    return payload.decode(message.get_content_charset() or "utf-8", errors="replace").strip()


def read_inputs(input_dir: Path) -> tuple[list[Starter], dict, dict, str]:
    for name, expected in INPUT_HASHES.items():
        source = input_dir / name
        if not source.is_file() or sha256(source) != expected:
            raise ValueError(f"输入 {name} 缺失或 SHA256 与冻结版本不一致")
    context = json.loads((input_dir / "Workflow_Context.json").read_text(encoding="utf-8"))
    if context != {"meta": {"schema_version": "0.1.0", "current_time": FIXED_ISO_TIME}}:
        raise ValueError("冻结运行时间已变化")
    channel = json.loads((input_dir / "IT_Channel.json").read_text(encoding="utf-8"))
    expected_channel = {"channels": [{"id": "C_IT_5107", "name": "it-general", "is_private": False}], "messages": []}
    if channel != expected_channel:
        raise ValueError("IT Slack 频道快照已变化")

    workbook = load_workbook(input_dir / "New_Hire_Onboarding_Tracker.xlsx", read_only=True, data_only=True)
    if workbook.sheetnames != ["April Starters"]:
        raise ValueError("入职工作簿结构已变化")
    starters: list[Starter] = []
    for row in workbook["April Starters"].iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        starters.append(Starter(
            source_row=int(row[0]), name=str(row[1]), employee_id=str(row[2]), department=str(row[3]),
            start_date=str(row[4]), location=str(row[5]), employment_type=str(row[6]), notes=str(row[7] or ""),
        ))
    workbook.close()
    frozen = [
        (2, "Ada Petrova", "EMP-7701", "Engineering", "2026-04-21", "Office - Austin", "Full-time", ""),
        (3, "Brian Osei", "EMP-7715", "Sales", "2026-04-21", "Remote - Chicago", "Full-time", ""),
        (4, "Carmen Delgado", "CON-2201", "Design", "2026-04-21", "Office - Austin", "Contractor", ""),
        (5, "David Nguyen", "EMP-7729", "Engineering", "2026-04-21", "Remote - Portland", "Full-time", ""),
        (6, "Elena Rossi", "EMP-7738", "Marketing", "2026-04-21", "Office - Austin", "Full-time", "Equipment already shipped from previous office transfer"),
        (7, "Femi Adeyemi", "CON-2215", "Engineering", "2026-04-21", "Remote - Lagos", "Independent Contractor", ""),
        (8, "Greta Holm", "EMP-7752", "Support", "2026-04-21", "Office - Austin", "Full-time", ""),
        (9, "Hugo Fernandez", "EMP-7766", "Engineering", "2026-04-21", "Remote - Madrid", "Full-time", ""),
    ]
    actual = [(s.source_row, s.name, s.employee_id, s.department, s.start_date, s.location, s.employment_type, s.notes) for s in starters]
    if actual != frozen:
        raise ValueError("八条冻结入职记录已变化")

    box = mailbox.mbox(input_dir / "Provisioning_Policy_Mailbox.mbox", create=False)
    messages = list(box)
    box.close()
    if len(messages) != 1:
        raise ValueError("配置政策邮箱必须恰好包含一封邮件")
    message = messages[0]
    body = decode_message(message)
    message_id = str(message.get("X-AutomationBench-Message-ID", ""))
    required = (
        "Full-time employees receive standard equipment package",
        "Contractors and consultants must provide their own equipment",
        "Remote employees get an additional monitor",
        "wireless keyboard, mouse, headset", "MacBook Pro 16-inch M4", "MacBook Air 15-inch M4",
    )
    if message_id != "msg_provision_policy_5107" or any(token not in body for token in required):
        raise ValueError("冻结配置政策证据不完整")
    policy = {
        "message_id": message_id, "thread_id": str(message.get("X-AutomationBench-Thread-ID", "")),
        "from": str(message.get("From", "")), "to": str(message.get("To", "")),
        "subject": str(message.get("Subject", "")), "date": str(message.get("Date", "")), "body": body,
    }
    return starters, policy, channel, context["meta"]["current_time"]


def derive_decisions(starters: list[Starter]) -> list[Decision]:
    decisions: list[Decision] = []
    for item in starters:
        contractor = "contractor" in item.employment_type.casefold() or "consultant" in item.employment_type.casefold()
        already = "already shipped" in item.notes.casefold()
        if contractor:
            status = "EXCLUDED_CONTRACTOR"
        elif already:
            status = "ALREADY_PROVISIONED"
        else:
            status = "PROVISION"
        if status == "PROVISION":
            laptop = "MacBook Pro 16-inch M4" if item.department == "Engineering" else "MacBook Air 15-inch M4"
            monitor = "YES" if item.location.startswith("Remote -") else "NO"
            keyboard = mouse = headset = "YES"
            jira_plan, email_list = "DRAFT_NOT_CREATED", "INCLUDE"
        else:
            laptop = "NONE"
            monitor = keyboard = mouse = headset = "NO"
            jira_plan, email_list = "NO_DRAFT", "EXCLUDE"
        decisions.append(Decision(item, contractor, already, status, laptop, monitor, keyboard, mouse, headset, jira_plan, email_list))
    expected = {
        "Ada Petrova": ("PROVISION", "MacBook Pro 16-inch M4", "NO"),
        "Brian Osei": ("PROVISION", "MacBook Air 15-inch M4", "YES"),
        "Carmen Delgado": ("EXCLUDED_CONTRACTOR", "NONE", "NO"),
        "David Nguyen": ("PROVISION", "MacBook Pro 16-inch M4", "YES"),
        "Elena Rossi": ("ALREADY_PROVISIONED", "NONE", "NO"),
        "Femi Adeyemi": ("EXCLUDED_CONTRACTOR", "NONE", "NO"),
        "Greta Holm": ("PROVISION", "MacBook Air 15-inch M4", "NO"),
        "Hugo Fernandez": ("PROVISION", "MacBook Pro 16-inch M4", "YES"),
    }
    if {d.source.name: (d.status, d.laptop, d.monitor) for d in decisions} != expected:
        raise ValueError("逐人配置真值与冻结任务不一致")
    return decisions


def add_table(sheet, name: str) -> None:
    table = Table(displayName=name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)


def style_sheet(sheet, widths: dict[str, float], footer: str, *, a3: bool = False) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", bold=True, color=WHITE, size=9)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)
    sheet.row_dimensions[1].height = 30
    grid = Side(style="hair", color="D9E2F3")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=8.5)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=grid, right=grid, bottom=grid)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if a3 else sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.oddFooter.center.text = footer
    sheet.print_area = sheet.dimensions


def normalize_ooxml(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="equipment-5107-ooxml-") as temporary:
        target = Path(temporary) / path.name
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as destination:
            axis_ids: dict[str, str] = {}
            for name in sorted(source.namelist()):
                data = source.read(name)
                if name == "docProps/core.xml":
                    try:
                        root = etree.fromstring(data)
                        ns = {"dcterms": "http://purl.org/dc/terms/", "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"}
                        for tag in ("created", "modified"):
                            for node in root.xpath(f"//dcterms:{tag}", namespaces=ns):
                                node.text = "2026-04-15T09:00:00Z"
                        for tag in ("lastModifiedBy", "revision"):
                            for node in root.xpath(f"//cp:{tag}", namespaces=ns):
                                node.text = "EvoCowork" if tag == "lastModifiedBy" else "1"
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except etree.XMLSyntaxError:
                        pass
                if name.startswith("xl/charts/chart") and name.endswith(".xml"):
                    try:
                        root = etree.fromstring(data)
                        ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
                        for node in root.xpath("//c:axId", namespaces=ns):
                            old = node.get("val")
                            if old and old not in axis_ids:
                                axis_ids[old] = str(51_070_001 + len(axis_ids))
                        for node in root.xpath("//c:axId | //c:crossAx", namespaces=ns):
                            old = node.get("val")
                            if old in axis_ids:
                                node.set("val", axis_ids[old])
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except etree.XMLSyntaxError:
                        pass
                info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                destination.writestr(info, data)
        shutil.copy2(target, path)


def recalculate_xlsx(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="equipment-5107-recalc-") as temporary:
        root = Path(temporary)
        source_dir, output_dir, profile_dir = root / "source", root / "output", root / "profile"
        source_dir.mkdir()
        output_dir.mkdir()
        source = source_dir / path.name
        shutil.copy2(path, source)
        process = subprocess.run([
            "libreoffice", "--headless", f"-env:UserInstallation=file://{profile_dir}",
            "--convert-to", "xlsx", "--outdir", str(output_dir), str(source),
        ], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300)
        converted = output_dir / path.name
        if process.returncode != 0 or not converted.is_file():
            raise RuntimeError(f"LibreOffice 公式重算失败：{process.stdout[-1000:]}")
        shutil.copy2(converted, path)


def write_workbook(path: Path, starters: list[Starter], decisions: list[Decision], policy_mail: dict, channel: dict, frozen_time: str) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.creator = "EvoCowork AutomationBench 5107 adaptation"
    workbook.title = "April 2026 Equipment Provisioning Audit"
    workbook.subject = "Formula-driven onboarding equipment control"
    workbook.description = "All people, devices, locations and policy rules trace to frozen AutomationBench 5107 inputs."
    fixed_dt = datetime(2026, 4, 15, 9, 0, 0)
    workbook.created = fixed_dt
    workbook.modified = fixed_dt

    source = workbook.create_sheet("Source_Starters")
    source.append(["Source Row", "Name", "Employee ID", "Department", "Start Date", "Location", "Employment Type", "Notes"])
    for item in starters:
        source.append([item.source_row, item.name, item.employee_id, item.department, item.start_date, item.location, item.employment_type, item.notes])
    add_table(source, "EquipmentSourceStarters")
    style_sheet(source, {"A": 11, "B": 21, "C": 16, "D": 19, "E": 16, "F": 24, "G": 24, "H": 55}, "Frozen source rows · AutomationBench 5107", a3=True)
    for row in range(2, 10):
        source.row_dimensions[row].height = 30
    source.row_dimensions[6].height = 44

    policy = workbook.create_sheet("Policy_Rules")
    policy.append(["Policy ID", "Rule", "Required Action", "Source Trace"])
    policy_rows = [
        ("P1", "Full-time employees", "Receive the standard equipment package", "Provisioning_Policy_Mailbox.mbox:msg_provision_policy_5107"),
        ("P2", "Contractors and consultants", "Provide their own equipment; no order or IT ticket", "Provisioning_Policy_Mailbox.mbox:msg_provision_policy_5107"),
        ("P3", "Remote employees", "Add one monitor", "Provisioning_Policy_Mailbox.mbox:msg_provision_policy_5107"),
        ("P4", "Standard package", "Laptop, wireless keyboard, mouse and headset", "Provisioning_Policy_Mailbox.mbox:msg_provision_policy_5107"),
        ("P5", "Department laptop", "Engineering: MacBook Pro 16-inch M4; all others: MacBook Air 15-inch M4", "Provisioning_Policy_Mailbox.mbox:msg_provision_policy_5107"),
    ]
    for row in policy_rows:
        policy.append(row)
    add_table(policy, "EquipmentPolicyRules")
    style_sheet(policy, {"A": 13, "B": 34, "C": 82, "D": 65}, "Five frozen provisioning rules", a3=True)
    for row in range(2, 7):
        policy.row_dimensions[row].height = 44

    evidence = workbook.create_sheet("Evidence_Register")
    evidence.append(["Record", "Evidence Type", "Evidence Ref", "Evidence Detail", "Frozen At", "Source Trace"])
    for item in starters:
        evidence.append([
            f"Source row {item.source_row} · {item.name}", "SOURCE_ROW", f"April Starters row {item.source_row}",
            f"{item.employee_id}; {item.department}; {item.start_date}; {item.location}; {item.employment_type}; notes {item.notes or '[blank]'}",
            frozen_time, f"New_Hire_Onboarding_Tracker.xlsx:April Starters row {item.source_row}",
        ])
    evidence.append(["Provisioning policy", "FULL_EMAIL", policy_mail["message_id"], policy_mail["body"], policy_mail["date"], "Provisioning_Policy_Mailbox.mbox:msg_provision_policy_5107"])
    evidence.append(["Workflow context", "FROZEN_TIME", "meta.current_time", frozen_time, frozen_time, "Workflow_Context.json:meta.current_time"])
    evidence.append(["IT channel", "CHANNEL_SNAPSHOT", channel["channels"][0]["id"], json.dumps(channel["channels"][0], ensure_ascii=False, sort_keys=True), frozen_time, "IT_Channel.json:channels[0]"])
    add_table(evidence, "EquipmentEvidenceRegister")
    style_sheet(evidence, {"A": 31, "B": 22, "C": 34, "D": 95, "E": 24, "F": 70}, "Evidence register · rows, policy, time and channel", a3=True)
    for row in range(2, 10):
        evidence.row_dimensions[row].height = 34
    evidence.row_dimensions[10].height = 100
    evidence.row_dimensions[11].height = 28
    evidence.row_dimensions[12].height = 34

    sheet = workbook.create_sheet("Provisioning_Decisions")
    sheet.append([
        "Source Row", "Name", "Employee ID", "Department", "Start Date", "Location", "Employment Type", "Notes",
        "Contractor", "Already Provisioned", "Final Status", "Laptop", "Additional Monitor",
        "Wireless Keyboard", "Mouse", "Headset", "Jira Plan", "IT Email List", "Source Trace",
    ])
    for row in range(2, 10):
        sheet.append([
            f"=Source_Starters!A{row}", f"=Source_Starters!B{row}", f"=Source_Starters!C{row}", f"=Source_Starters!D{row}",
            f"=Source_Starters!E{row}", f"=Source_Starters!F{row}", f"=Source_Starters!G{row}", f'=IF(Source_Starters!H{row}="","",Source_Starters!H{row})',
            f'=IF(OR(ISNUMBER(SEARCH("Contractor",G{row})),ISNUMBER(SEARCH("Consultant",G{row}))),"YES","NO")',
            f'=IF(ISNUMBER(SEARCH("already shipped",H{row})),"YES","NO")',
            f'=IF(I{row}="YES","EXCLUDED_CONTRACTOR",IF(J{row}="YES","ALREADY_PROVISIONED","PROVISION"))',
            f'=IF(K{row}<>"PROVISION","NONE",IF(D{row}="Engineering","MacBook Pro 16-inch M4","MacBook Air 15-inch M4"))',
            f'=IF(AND(K{row}="PROVISION",ISNUMBER(SEARCH("Remote -",F{row}))),"YES","NO")',
            f'=IF(K{row}="PROVISION","YES","NO")', f'=IF(K{row}="PROVISION","YES","NO")', f'=IF(K{row}="PROVISION","YES","NO")',
            f'=IF(K{row}="PROVISION","DRAFT_NOT_CREATED","NO_DRAFT")', f'=IF(K{row}="PROVISION","INCLUDE","EXCLUDE")',
            f"=Evidence_Register!F{row}",
        ])
    for status, color in (("PROVISION", GREEN), ("EXCLUDED_CONTRACTOR", RED), ("ALREADY_PROVISIONED", AMBER)):
        sheet.conditional_formatting.add("K2:K9", FormulaRule(formula=[f'K2="{status}"'], fill=PatternFill("solid", fgColor=color)))
    add_table(sheet, "EquipmentProvisioningDecisions")
    style_sheet(sheet, {
        "A": 10, "B": 19, "C": 15, "D": 18, "E": 15, "F": 23, "G": 23, "H": 45,
        "I": 14, "J": 19, "K": 24, "L": 28, "M": 19, "N": 18, "O": 11, "P": 12,
        "Q": 20, "R": 16, "S": 65,
    }, "Formula decisions · exclusion precedence, then department/location", a3=True)
    for row in range(2, 10):
        sheet.row_dimensions[row].height = 38
    sheet.row_dimensions[6].height = 50

    queue = workbook.create_sheet("Ticket_Queue")
    queue.append(["Draft ID", "Name", "Employee ID", "Department", "Location", "Start Date", "Laptop", "Monitor", "Keyboard", "Mouse", "Headset", "Status", "Source Trace"])
    source_rows = [2, 3, 5, 8, 9]
    for index, decision_row in enumerate(source_rows, start=1):
        queue.append([
            f"JIRA-{index:03d}", f"=Provisioning_Decisions!B{decision_row}", f"=Provisioning_Decisions!C{decision_row}",
            f"=Provisioning_Decisions!D{decision_row}", f"=Provisioning_Decisions!F{decision_row}", f"=Provisioning_Decisions!E{decision_row}",
            f"=Provisioning_Decisions!L{decision_row}", f"=Provisioning_Decisions!M{decision_row}", f"=Provisioning_Decisions!N{decision_row}",
            f"=Provisioning_Decisions!O{decision_row}", f"=Provisioning_Decisions!P{decision_row}", "DRAFT_NOT_CREATED", f"=Provisioning_Decisions!S{decision_row}",
        ])
    add_table(queue, "EquipmentTicketQueue")
    style_sheet(queue, {"A": 13, "B": 20, "C": 15, "D": 18, "E": 24, "F": 15, "G": 28, "H": 12, "I": 13, "J": 10, "K": 11, "L": 21, "M": 68}, "Five Jira drafts · zero created tickets", a3=True)
    for row in range(2, 7):
        queue.row_dimensions[row].height = 38

    summary = workbook.create_sheet("Control_Summary")
    summary.append(["Metric", "Value", "Control Meaning"])
    rows = [
        ("Total starters", "=COUNTA(Provisioning_Decisions!B2:B9)", "All frozen source rows"),
        ("Provision", '=COUNTIF(Provisioning_Decisions!K2:K9,"PROVISION")', "Five Jira drafts"),
        ("Excluded contractors", '=COUNTIF(Provisioning_Decisions!K2:K9,"EXCLUDED_CONTRACTOR")', "Carmen and Femi"),
        ("Already provisioned", '=COUNTIF(Provisioning_Decisions!K2:K9,"ALREADY_PROVISIONED")', "Elena; no duplicate"),
        ("MacBook Pro 16-inch M4", '=COUNTIF(Provisioning_Decisions!L2:L9,"MacBook Pro 16-inch M4")', "Engineering provision cases"),
        ("MacBook Air 15-inch M4", '=COUNTIF(Provisioning_Decisions!L2:L9,"MacBook Air 15-inch M4")', "Non-engineering provision cases"),
        ("Additional monitors", '=COUNTIF(Provisioning_Decisions!M2:M9,"YES")', "Remote provision cases"),
        ("Jira drafts", '=COUNTIF(Provisioning_Decisions!Q2:Q9,"DRAFT_NOT_CREATED")', "Exactly five"),
        ("Emails sent", "=0", "One draft only"),
        ("Orders placed", "=0", "No procurement execution"),
    ]
    for row in rows:
        summary.append(row)
    add_table(summary, "EquipmentControlSummary")
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Provisioning status distribution"
    chart.y_axis.title = "People"
    chart.x_axis.title = "Status"
    chart.legend = None
    chart.height = 7.0
    chart.width = 14.0
    chart.add_data(Reference(summary, min_col=2, min_row=3, max_row=5), titles_from_data=False)
    chart.set_categories(Reference(summary, min_col=1, min_row=3, max_row=5))
    summary.add_chart(chart, "E2")
    style_sheet(summary, {"A": 36, "B": 18, "C": 50}, "Control totals · formulas linked to Provisioning_Decisions")
    summary.print_area = "A1:L11"

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(path)
    recalculate_xlsx(path)
    normalize_ooxml(path)


def set_cell_shading(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def format_doc_table(table, widths: list[float] | None = None, font_size: float = 8.0) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for cell in table.rows[0].cells:
        set_cell_shading(cell, NAVY)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(font_size)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths and index < len(widths):
                cell.width = Cm(widths[index])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(font_size)


def add_title(document: Document, title: str, subtitle: str) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(23, 54, 93)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle)
    run.italic = True
    run.font.name = "Aptos"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(89, 89, 89)


def add_heading(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(31, 78, 121)


def add_body(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(9.0)


def write_brief(path: Path, decisions: list[Decision]) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.1)
    section.bottom_margin = Cm(1.1)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    core = document.core_properties
    core.title = "April 2026 Equipment Provisioning Control Brief"
    core.author = "EvoCowork AutomationBench 5107 adaptation"
    core.subject = "Eight-person onboarding equipment policy brief"
    core.created = datetime(2026, 4, 15, 9, 0, 0)
    core.modified = datetime(2026, 4, 15, 9, 0, 0)
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(9)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run("HR OPS / IT CONTROL · APRIL 2026")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Frozen 2026-04-15 · Drafts only · No ticket, message, order or shipment side effect")

    add_title(document, "Equipment Provisioning Control Brief", "April 2026 starters · policy-controlled offline package")
    add_heading(document, "Executive control result")
    add_body(document, "Eight source rows were reconciled against the CFO policy. Five people are PROVISION, two are EXCLUDED_CONTRACTOR and one is ALREADY_PROVISIONED. The package requires 3 MacBook Pro 16-inch M4, 2 MacBook Air 15-inch M4 and 3 additional monitors.")
    metrics = document.add_table(rows=2, cols=7)
    labels = ["Starters", "Provision", "Contractors", "Already shipped", "Pro", "Air", "Monitors"]
    values = ["8", "5", "2", "1", "3", "2", "3"]
    for index, label in enumerate(labels):
        metrics.cell(0, index).text = label
        metrics.cell(1, index).text = values[index]
    format_doc_table(metrics, font_size=8.2)
    add_heading(document, "Decision order")
    add_body(document, "Exclude contractors and consultants first; suppress duplicate work when Notes says equipment already shipped; then assign Engineering to Pro and all other provision cases to Air; add a monitor only for Remote locations; give every provision case a wireless keyboard, mouse and headset.")
    add_heading(document, "Evidence chain")
    add_body(document, "New_Hire_Onboarding_Tracker.xlsx supplies eight rows; Provisioning_Policy_Mailbox.mbox / msg_provision_policy_5107 supplies five rules; Workflow_Context.json freezes 2026-04-15T09:00:00Z; IT_Channel.json identifies C_IT_5107 / it-general.")
    add_heading(document, "Side-effect boundary")
    add_body(document, "Five Jira issues, one IT email and one channel message are drafts only. Purchase_Execution_Control.json records ordered=false, purchase_record_count=0 and jira_created_count=0. Nothing was created, sent, posted, ordered or shipped.")

    document.add_page_break()
    add_title(document, "Starter Decision Register", "Employment/note precedence -> department/location package")
    table = document.add_table(rows=1, cols=7)
    for index, value in enumerate(["Name", "ID", "Department", "Location", "Status", "Laptop", "Monitor"]):
        table.cell(0, index).text = value
    for item in decisions:
        cells = table.add_row().cells
        values = [item.source.name, item.source.employee_id, item.source.department, item.source.location, item.status, item.laptop, item.monitor]
        for index, value in enumerate(values):
            cells[index].text = value
    format_doc_table(table, widths=[3.7, 2.4, 3.1, 4.0, 4.3, 5.2, 1.8], font_size=7.5)
    add_heading(document, "Package and exclusion checks")
    add_body(document, "Ada, Brian, David, Greta and Hugo each receive wireless keyboard=YES, mouse=YES and headset=YES. Brian, David and Hugo are the only additional-monitor cases. Carmen and Femi are contractor exclusions. Elena is not duplicated because her equipment was already shipped from the previous office transfer.")
    facts = document.add_table(rows=1, cols=3)
    for index, value in enumerate(["Group", "Required handling", "Forbidden outcome"]):
        facts.cell(0, index).text = value
    rows = [
        ("Ada; David; Hugo", "Pro 16-inch M4; standard peripherals", "Air model; contractor exclusion"),
        ("Brian; Greta", "Air 15-inch M4; standard peripherals", "Pro model"),
        ("Brian; David; Hugo", "One additional monitor each", "Monitor for office-only cases"),
        ("Carmen; Femi", "EXCLUDED_CONTRACTOR", "Any Jira, email list or order"),
        ("Elena", "ALREADY_PROVISIONED", "Duplicate Jira, email list or order"),
    ]
    for row in rows:
        cells = facts.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    format_doc_table(facts, widths=[7.0, 9.0, 9.0], font_size=7.8)

    document.add_page_break()
    add_title(document, "Draft & Execution Control", "Five Jira drafts, one IT list, one aggregate channel draft, zero side effects")
    add_heading(document, "Jira draft register")
    jira = document.add_table(rows=1, cols=5)
    for index, value in enumerate(["Draft", "Person", "Laptop", "Monitor", "Status"]):
        jira.cell(0, index).text = value
    for sequence, item in enumerate([d for d in decisions if d.status == "PROVISION"], start=1):
        cells = jira.add_row().cells
        values = [f"JIRA-{sequence:03d}", item.source.name, item.laptop, item.monitor, "DRAFT_NOT_CREATED"]
        for index, value in enumerate(values):
            cells[index].text = value
    format_doc_table(jira, widths=[2.8, 4.5, 8.0, 2.5, 5.0], font_size=7.8)
    add_heading(document, "Five provisioning rules")
    policy = document.add_table(rows=1, cols=2)
    policy.cell(0, 0).text = "Rule"
    policy.cell(0, 1).text = "Controlled interpretation"
    rules = [
        ("Full-time", "Standard equipment package"),
        ("Contractor / consultant", "Own equipment; no order or IT ticket"),
        ("Remote", "One additional monitor"),
        ("Standard package", "Laptop + wireless keyboard + mouse + headset"),
        ("Department", "Engineering Pro 16-inch M4; all others Air 15-inch M4"),
    ]
    for row in rules:
        cells = policy.add_row().cells
        cells[0].text, cells[1].text = row
    format_doc_table(policy, widths=[7.0, 18.0], font_size=7.8)
    add_heading(document, "Privacy and final attestation")
    add_body(document, "The C_IT_5107 / it-general draft contains only aggregate 8/5/2/1 and 3/2/3 counts, with no name, Employee ID, location or per-person equipment. The IT email remains DRAFT_NOT_SENT; all Jira issues remain DRAFT_NOT_CREATED. ordered=false; purchase records=0; Jira created=0.")
    document.save(path)
    normalize_ooxml(path)


def write_jira(path: Path, decisions: list[Decision]) -> None:
    issues = []
    for sequence, item in enumerate([d for d in decisions if d.status == "PROVISION"], start=1):
        issues.append({
            "draft_id": f"JIRA-{sequence:03d}", "status": "DRAFT_NOT_CREATED", "created": False,
            "summary": f"Provision equipment for {item.source.name} ({item.source.employee_id})",
            "person": item.source.name, "employee_id": item.source.employee_id,
            "department": item.source.department, "location": item.source.location,
            "start_date": item.source.start_date, "laptop": item.laptop,
            "peripherals": {"wireless_keyboard": True, "mouse": True, "headset": True, "additional_monitor": item.monitor == "YES"},
            "source": {"file": "New_Hire_Onboarding_Tracker.xlsx", "sheet": "April Starters", "source_row": item.source.source_row},
            "policy_source": "Provisioning_Policy_Mailbox.mbox:msg_provision_policy_5107",
        })
    payload = {"status": "DRAFT_NOT_CREATED", "created": False, "issue_count": 5, "issues": issues, "created_issue_keys": []}
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_mbox(path: Path, decisions: list[Decision]) -> None:
    selected = [d for d in decisions if d.status == "PROVISION"]
    lines = []
    for item in selected:
        lines.append(
            f"- {item.source.name} | {item.source.employee_id} | {item.source.department} | {item.source.location} | start {item.source.start_date} | "
            f"{item.laptop} | wireless keyboard | mouse | headset | additional monitor {item.monitor}"
        )
    message = EmailMessage(policy=SMTP)
    message["From"] = "hr-ops@company.example.com"
    message["To"] = "it-ops@company.example.com"
    message["Subject"] = "DRAFT — April 21 starter equipment list"
    message["Date"] = "Wed, 15 Apr 2026 09:00:00 +0000"
    message["Message-ID"] = "<automationbench-5107-it-list@evocowork.invalid>"
    message["X-EvoCowork-Status"] = "DRAFT_NOT_SENT"
    message["X-EvoCowork-Source"] = "AutomationBench-5107"
    body = (
        "DRAFT_NOT_SENT\n\nIT Operations,\n\n"
        "The controlled equipment list for eligible April 21 starters is:\n"
        + "\n".join(lines)
        + "\n\nFive Jira drafts have been prepared. Three source rows are excluded by employment type or prior shipment. This email has not been sent. No ticket was created, no equipment was ordered and no shipment is claimed."
    )
    message.set_content(body + "\n", charset="utf-8")
    buffer = BytesIO()
    BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
    path.write_bytes(b"From hr-ops@company.example.com Wed Apr 15 09:00:00 2026\n" + buffer.getvalue().replace(b"\r\n", b"\n") + b"\n")


def write_slack(path: Path) -> None:
    payload = {
        "source_channel": {"id": "C_IT_5107", "name": "it-general", "is_private": False},
        "status": "DRAFT_NOT_POSTED", "posted": False,
        "summary": {"total_starters": 8, "provision": 5, "excluded_contractors": 2, "already_provisioned": 1, "macbook_pro": 3, "macbook_air": 2, "additional_monitors": 3},
        "draft_message": {"text": "April 21 equipment control: 8 starters reviewed; 5 PROVISION, 2 EXCLUDED_CONTRACTOR, 1 ALREADY_PROVISIONED; 3 Pro, 2 Air and 3 additional monitors. Drafts prepared; no ticket, message or order executed."},
        "privacy_control": {"aggregate_only": True, "names_included": False, "employee_ids_included": False, "locations_included": False, "per_person_equipment_included": False},
        "source_trace": ["Equipment_Provisioning_Audit.xlsx:Control_Summary", "IT_Channel.json:channels[0]"],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_control(path: Path) -> None:
    payload = {
        "decision": "DRAFT_ONLY_NO_ORDER", "status": "DRAFT_ONLY_NO_ORDER", "ordered": False,
        "purchase_record_count": 0, "jira_created_count": 0, "email_sent_count": 0, "slack_posted_count": 0,
        "order_ids": [], "created_issue_keys": [], "shipment_ids": [],
        "draft_counts": {"jira": 5, "email": 1, "slack": 1},
        "control_reason": "The benchmark prepares offline routing artifacts and does not execute Jira, email, Slack, procurement or shipping side effects.",
        "source_trace": ["New_Hire_Onboarding_Tracker.xlsx:April Starters", "Provisioning_Policy_Mailbox.mbox:msg_provision_policy_5107", "Equipment_Provisioning_Audit.xlsx:Ticket_Queue"],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=Path("/app/input_files"))
    parser.add_argument("--output-dir", type=Path, default=Path("/app/output"))
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    if any(args.output_dir.iterdir()):
        raise ValueError("输出目录必须为空，避免遗留文件污染交付")
    starters, policy_mail, channel, frozen_time = read_inputs(args.input_dir)
    decisions = derive_decisions(starters)
    write_workbook(args.output_dir / XLSX_NAME, starters, decisions, policy_mail, channel, frozen_time)
    write_brief(args.output_dir / DOCX_NAME, decisions)
    write_jira(args.output_dir / JIRA_NAME, decisions)
    write_mbox(args.output_dir / MBOX_NAME, decisions)
    write_slack(args.output_dir / SLACK_NAME)
    write_control(args.output_dir / CONTROL_NAME)


if __name__ == "__main__":
    main()
