#!/usr/bin/env python3
"""AutomationBench 5107 入职设备配置题的逐题确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import mailbox
import re
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from email.utils import getaddresses
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook


TASK_ID = "cowork-evo/automationbench-5107-equipment-provisioning-package"
XLSX_NAME = "Equipment_Provisioning_Audit.xlsx"
DOCX_NAME = "Provisioning_Control_Brief.docx"
JIRA_NAME = "Provisioning_Jira_Drafts.json"
MBOX_NAME = "IT_Equipment_List_Draft.mbox"
SLACK_NAME = "IT_Channel_Draft.json"
CONTROL_NAME = "Purchase_Execution_Control.json"
REQUIRED_FILES = {XLSX_NAME, DOCX_NAME, JIRA_NAME, MBOX_NAME, SLACK_NAME, CONTROL_NAME}
INPUT_HASHES = {
    "New_Hire_Onboarding_Tracker.xlsx": "aa8ac492cf0995a480de8ff5a9435c19b601a457ad4d358f029c9b0525e0d20d",
    "Provisioning_Policy_Mailbox.mbox": "a9b0849b12b6b4642a7ffea9227ba39fe2460ea43435bd33a4f07fb7137e6124",
    "IT_Channel.json": "c98d2dbebd2d0831d2b82e5309bd71fbfd7cb3bc6e718fcc502d5d27432fab32",
    "Workflow_Context.json": "76c4a0657bc27f2fd830b10613ec7ceae3adcdc5798aeee0f822d27bf69bbc99",
}

WEIGHTS = {
    "R001": 4, "R002": 6, "R003": 6, "R004": 7,
    "R005": 4, "R006": 4, "R007": 5, "R008": 4,
    "R009": 5, "R010": 5, "R011": 4, "R012": 4,
    "R013": 8, "R014": 8, "R015": 5, "R016": 6,
    "R017": 6, "R018": 5, "R019": 4,
}
HURDLE_CAPS = {
    "R001": 0.0, "R002": 0.2, "R003": 0.2, "R004": 0.3,
    "R005": 0.35, "R006": 0.35, "R007": 0.25, "R008": 0.35,
    "R009": 0.25, "R010": 0.25, "R011": 0.35, "R012": 0.35,
    "R013": 0.2, "R014": 0.2, "R015": 0.25, "R016": 0.2,
    "R017": 0.35, "R018": 0.3,
}

FROZEN_STARTERS = [
    (2, "Ada Petrova", "EMP-7701", "Engineering", "2026-04-21", "Office - Austin", "Full-time", ""),
    (3, "Brian Osei", "EMP-7715", "Sales", "2026-04-21", "Remote - Chicago", "Full-time", ""),
    (4, "Carmen Delgado", "CON-2201", "Design", "2026-04-21", "Office - Austin", "Contractor", ""),
    (5, "David Nguyen", "EMP-7729", "Engineering", "2026-04-21", "Remote - Portland", "Full-time", ""),
    (6, "Elena Rossi", "EMP-7738", "Marketing", "2026-04-21", "Office - Austin", "Full-time", "Equipment already shipped from previous office transfer"),
    (7, "Femi Adeyemi", "CON-2215", "Engineering", "2026-04-21", "Remote - Lagos", "Independent Contractor", ""),
    (8, "Greta Holm", "EMP-7752", "Support", "2026-04-21", "Office - Austin", "Full-time", ""),
    (9, "Hugo Fernandez", "EMP-7766", "Engineering", "2026-04-21", "Remote - Madrid", "Full-time", ""),
]
EXPECTED = {
    "Ada Petrova": ("NO", "NO", "PROVISION", "MacBook Pro 16-inch M4", "NO", "YES", "YES", "YES", "DRAFT_NOT_CREATED", "INCLUDE"),
    "Brian Osei": ("NO", "NO", "PROVISION", "MacBook Air 15-inch M4", "YES", "YES", "YES", "YES", "DRAFT_NOT_CREATED", "INCLUDE"),
    "Carmen Delgado": ("YES", "NO", "EXCLUDED_CONTRACTOR", "NONE", "NO", "NO", "NO", "NO", "NO_DRAFT", "EXCLUDE"),
    "David Nguyen": ("NO", "NO", "PROVISION", "MacBook Pro 16-inch M4", "YES", "YES", "YES", "YES", "DRAFT_NOT_CREATED", "INCLUDE"),
    "Elena Rossi": ("NO", "YES", "ALREADY_PROVISIONED", "NONE", "NO", "NO", "NO", "NO", "NO_DRAFT", "EXCLUDE"),
    "Femi Adeyemi": ("YES", "NO", "EXCLUDED_CONTRACTOR", "NONE", "NO", "NO", "NO", "NO", "NO_DRAFT", "EXCLUDE"),
    "Greta Holm": ("NO", "NO", "PROVISION", "MacBook Air 15-inch M4", "NO", "YES", "YES", "YES", "DRAFT_NOT_CREATED", "INCLUDE"),
    "Hugo Fernandez": ("NO", "NO", "PROVISION", "MacBook Pro 16-inch M4", "YES", "YES", "YES", "YES", "DRAFT_NOT_CREATED", "INCLUDE"),
}
CRITERION_BY_NAME = {
    "Ada Petrova": "R005", "Brian Osei": "R006", "Carmen Delgado": "R007",
    "David Nguyen": "R008", "Elena Rossi": "R009", "Femi Adeyemi": "R010",
    "Greta Holm": "R011", "Hugo Fernandez": "R012",
}
PROVISION_NAMES = ("Ada Petrova", "Brian Osei", "David Nguyen", "Greta Holm", "Hugo Fernandez")
EXCLUDED_NAMES = ("Carmen Delgado", "Elena Rossi", "Femi Adeyemi")


@dataclass
class CheckLedger:
    values: dict[str, dict[str, Any]]

    def __init__(self) -> None:
        self.values = {key: {"passed": True, "details": []} for key in WEIGHTS}

    def require(self, criterion: str, condition: bool, detail: str) -> None:
        if not condition:
            self.values[criterion]["passed"] = False
            self.values[criterion]["details"].append(detail)

    def fail(self, criterion: str, detail: str) -> None:
        self.require(criterion, False, detail)

    def passed(self, criterion: str) -> bool:
        return bool(self.values[criterion]["passed"])


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def task_root() -> Path:
    return Path(__file__).resolve().parents[1]


def input_root() -> Path:
    container = Path("/app/input_files")
    return container if container.is_dir() else task_root() / "data" / "input_files"


def normalize_scalar(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def value_rows(sheet, start: int, end: int, columns: int) -> list[tuple[Any, ...]]:
    return [tuple(normalize_scalar(sheet.cell(row, column).value) for column in range(1, columns + 1)) for row in range(start, end + 1)]


def decode_message(message) -> str:
    if message.is_multipart():
        parts: list[str] = []
        for part in message.walk():
            if part.get_content_type() == "text/plain":
                payload = part.get_payload(decode=True) or b""
                parts.append(payload.decode(part.get_content_charset() or "utf-8", errors="replace"))
        return "\n".join(parts).strip()
    payload = message.get_payload(decode=True) or b""
    return payload.decode(message.get_content_charset() or "utf-8", errors="replace").strip()


def addresses(message, header: str) -> set[str]:
    return {address.casefold() for _, address in getaddresses(message.get_all(header, [])) if address}


def render_office(path: Path) -> tuple[int, str]:
    with tempfile.TemporaryDirectory(prefix="judge-5107-render-") as temporary:
        root = Path(temporary)
        output = root / "output"
        profile = root / "profile"
        output.mkdir()
        process = subprocess.run(
            ["libreoffice", "--headless", f"-env:UserInstallation=file://{profile}", "--convert-to", "pdf", "--outdir", str(output), str(path)],
            stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300,
        )
        pdf = output / f"{path.stem}.pdf"
        if process.returncode != 0 or not pdf.is_file():
            raise RuntimeError(process.stdout[-1000:])
        info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True, timeout=60)
        match = re.search(r"^Pages:\s+(\d+)", info.stdout, flags=re.MULTILINE)
        if info.returncode != 0 or not match:
            raise RuntimeError("无法读取 PDF 页数")
        text = subprocess.run(["pdftotext", str(pdf), "-"], capture_output=True, text=True, timeout=60)
        if text.returncode != 0:
            raise RuntimeError("无法提取 PDF 文本")
        return int(match.group(1)), text.stdout


def formula_depends_on(value: Any, *tokens: str) -> bool:
    formula = str(value or "")
    return formula.startswith("=") and all(token.casefold() in formula.casefold() for token in tokens)


def workbook_checks(artifact_dir: Path, checks: CheckLedger) -> dict[str, dict[str, Any]]:
    formula_book = load_workbook(artifact_dir / XLSX_NAME, data_only=False)
    value_book = load_workbook(artifact_dir / XLSX_NAME, data_only=True)
    expected_sheets = ["Source_Starters", "Policy_Rules", "Evidence_Register", "Provisioning_Decisions", "Ticket_Queue", "Control_Summary"]
    checks.require("R001", formula_book.sheetnames == expected_sheets, f"Excel 工作表应为 {expected_sheets}")

    source_rows = value_rows(value_book["Source_Starters"], 2, 9, 8)
    checks.require("R002", source_rows == FROZEN_STARTERS, "Source_Starters 未逐行保留八条冻结来源")

    policy_rows = value_rows(value_book["Policy_Rules"], 2, 6, 4)
    expected_policy = [
        ("P1", "Full-time employees", "Receive the standard equipment package", "Provisioning_Policy_Mailbox.mbox:msg_provision_policy_5107"),
        ("P2", "Contractors and consultants", "Provide their own equipment; no order or IT ticket", "Provisioning_Policy_Mailbox.mbox:msg_provision_policy_5107"),
        ("P3", "Remote employees", "Add one monitor", "Provisioning_Policy_Mailbox.mbox:msg_provision_policy_5107"),
        ("P4", "Standard package", "Laptop, wireless keyboard, mouse and headset", "Provisioning_Policy_Mailbox.mbox:msg_provision_policy_5107"),
        ("P5", "Department laptop", "Engineering: MacBook Pro 16-inch M4; all others: MacBook Air 15-inch M4", "Provisioning_Policy_Mailbox.mbox:msg_provision_policy_5107"),
    ]
    checks.require("R003", policy_rows == expected_policy, "Policy_Rules 的五项冻结政策不准确")

    evidence_rows = value_rows(value_book["Evidence_Register"], 2, 12, 6)
    for index, frozen in enumerate(FROZEN_STARTERS):
        row = evidence_rows[index]
        checks.require("R003", frozen[1] in str(row[0]) and f"row {frozen[0]}" in str(row[2]), f"{frozen[1]} 证据记录错位")
        for token in frozen[2:7]:
            checks.require("R003", str(token) in str(row[3]), f"{frozen[1]} 证据缺少 {token}")
        checks.require("R018", str(row[5]).endswith(f"row {frozen[0]}"), f"{frozen[1]} 来源追踪错误")
    policy_evidence, time_evidence, channel_evidence = evidence_rows[8:11]
    checks.require("R003", policy_evidence[1:3] == ("FULL_EMAIL", "msg_provision_policy_5107"), "完整政策邮件证据缺失")
    for token in ("Full-time employees", "Contractors and consultants", "Remote employees", "wireless keyboard, mouse, headset", "MacBook Pro 16-inch M4", "MacBook Air 15-inch M4"):
        checks.require("R003", token in str(policy_evidence[3]), f"政策邮件证据缺少 {token}")
    checks.require("R003", time_evidence[1:4] == ("FROZEN_TIME", "meta.current_time", "2026-04-15T09:00:00Z"), "冻结时间证据错误")
    checks.require("R003", channel_evidence[1:3] == ("CHANNEL_SNAPSHOT", "C_IT_5107") and '"name": "it-general"' in str(channel_evidence[3]), "IT 频道证据错误")

    formula_sheet = formula_book["Provisioning_Decisions"]
    value_sheet = value_book["Provisioning_Decisions"]
    for row in range(2, 10):
        for column in range(1, 8):
            expected_ref = f"=Source_Starters!{chr(64 + column)}{row}"
            checks.require("R004", str(formula_sheet.cell(row, column).value) == expected_ref, f"决策表 {row},{column} 未链接来源表")
        checks.require("R004", formula_depends_on(formula_sheet.cell(row, 8).value, f"Source_Starters!H{row}"), f"决策表 {row},8 未以空值安全方式链接备注")
        dependencies = {
            9: (f"G{row}",), 10: (f"H{row}",), 11: (f"I{row}", f"J{row}"),
            12: (f"K{row}", f"D{row}"), 13: (f"K{row}", f"F{row}"),
            14: (f"K{row}",), 15: (f"K{row}",), 16: (f"K{row}",),
            17: (f"K{row}",), 18: (f"K{row}",), 19: (f"Evidence_Register!F{row}",),
        }
        for column, tokens in dependencies.items():
            checks.require("R004", formula_depends_on(formula_sheet.cell(row, column).value, *tokens), f"决策表 {row},{column} 公式未依赖所需证据")
    summary_formulas = [formula_book["Control_Summary"].cell(row, 2).value for row in range(2, 12)]
    checks.require("R004", all(str(value).startswith("=") for value in summary_formulas), "Control_Summary 存在静态化汇总值")
    checks.require("R004", all("Provisioning_Decisions" in str(value) for value in summary_formulas[:8]) and summary_formulas[8:] == ["=0", "=0"], "汇总公式未链接决策表或零执行边界错误")
    queue_formula = formula_book["Ticket_Queue"]
    queue_formula_ok = True
    for row in range(2, 7):
        queue_formula_ok = queue_formula_ok and all(formula_depends_on(queue_formula.cell(row, column).value, "Provisioning_Decisions") for column in list(range(2, 12)) + [13])
    checks.require("R004", queue_formula_ok, "Ticket_Queue 存在脱离决策表的静态字段")

    decisions: dict[str, dict[str, Any]] = {}
    for row in range(2, 10):
        values = [normalize_scalar(value_sheet.cell(row, column).value) for column in range(1, 20)]
        decisions[str(values[1])] = {
            "source_row": values[0], "employee_id": values[2], "department": values[3], "start_date": values[4],
            "location": values[5], "employment_type": values[6], "notes": values[7], "contractor": values[8],
            "already": values[9], "status": values[10], "laptop": values[11], "monitor": values[12],
            "keyboard": values[13], "mouse": values[14], "headset": values[15], "jira": values[16],
            "email": values[17], "trace": values[18],
        }
    checks.require("R002", set(decisions) == set(EXPECTED), "决策表人员集合错误")
    frozen_by_name = {row[1]: row for row in FROZEN_STARTERS}
    for name, expected in EXPECTED.items():
        actual = decisions.get(name, {})
        observed = tuple(actual.get(key) for key in ("contractor", "already", "status", "laptop", "monitor", "keyboard", "mouse", "headset", "jira", "email"))
        criterion = CRITERION_BY_NAME[name]
        checks.require(criterion, observed == expected, f"{name} 配置/排除错误：{observed}")
        frozen = frozen_by_name[name]
        checks.require("R002", tuple(actual.get(key) for key in ("source_row", "employee_id", "department", "start_date", "location", "employment_type", "notes")) == (frozen[0], frozen[2], frozen[3], frozen[4], frozen[5], frozen[6], frozen[7]), f"{name} 来源字段变形")
        checks.require("R018", actual.get("trace") == f"New_Hire_Onboarding_Tracker.xlsx:April Starters row {frozen[0]}", f"{name} 决策追踪错误")

    summary_values = [normalize_scalar(value_book["Control_Summary"].cell(row, 2).value) for row in range(2, 12)]
    checks.require("R004", summary_values == [8, 5, 2, 1, 3, 2, 3, 5, 0, 0], f"汇总值错误：{summary_values}")
    queue_rows = value_rows(value_book["Ticket_Queue"], 2, 6, 13)
    expected_queue_names = list(PROVISION_NAMES)
    checks.require("R013", [row[0] for row in queue_rows] == [f"JIRA-{index:03d}" for index in range(1, 6)] and [row[1] for row in queue_rows] == expected_queue_names, "Ticket_Queue 不是五个合格人员草稿")
    for row, name in zip(queue_rows, expected_queue_names):
        frozen = frozen_by_name[name]
        expected = EXPECTED[name]
        checks.require(CRITERION_BY_NAME[name], row[2:12] == (frozen[2], frozen[3], frozen[5], frozen[4], expected[3], expected[4], "YES", "YES", "YES", "DRAFT_NOT_CREATED"), f"{name} Ticket_Queue 配置错误")
        checks.require("R018", row[12] == f"New_Hire_Onboarding_Tracker.xlsx:April Starters row {frozen[0]}", f"{name} Ticket_Queue 追踪错误")

    total_tables = sum(len(sheet.tables) for sheet in formula_book.worksheets)
    total_cf = sum(len(sheet.conditional_formatting) for sheet in formula_book.worksheets)
    total_charts = sum(len(sheet._charts) for sheet in formula_book.worksheets)
    checks.require("R019", total_tables == 6, f"应有6张原生表，实际{total_tables}")
    checks.require("R019", total_cf >= 1, "缺少决策条件格式")
    checks.require("R019", total_charts >= 1, "缺少可编辑原生图表")
    with zipfile.ZipFile(artifact_dir / XLSX_NAME) as archive:
        names = archive.namelist()
        checks.require("R001", not any(name.endswith("vbaProject.bin") for name in names), "工作簿包含宏")
        checks.require("R001", not any(name.startswith("xl/externalLinks/") for name in names), "工作簿包含外部链接")
    formula_book.close()
    value_book.close()
    return decisions


def jira_checks(artifact_dir: Path, checks: CheckLedger) -> dict[str, dict[str, Any]]:
    payload = json.loads((artifact_dir / JIRA_NAME).read_text(encoding="utf-8"))
    checks.require("R013", payload.get("status") == "DRAFT_NOT_CREATED" and payload.get("created") is False, "Jira 顶层状态不是未创建草稿")
    checks.require("R013", payload.get("issue_count") == 5 and payload.get("created_issue_keys") == [], "Jira 草稿数或 created issue key 错误")
    issues = payload.get("issues")
    checks.require("R013", isinstance(issues, list) and len(issues) == 5, "Jira 必须恰好有5个独立草稿")
    if not isinstance(issues, list):
        return {}
    by_name = {str(issue.get("person")): issue for issue in issues if isinstance(issue, dict)}
    checks.require("R013", set(by_name) == set(PROVISION_NAMES), "Jira 人员集合错误或包含排除人员")
    frozen_by_name = {row[1]: row for row in FROZEN_STARTERS}
    expected_ids = {name: f"JIRA-{index:03d}" for index, name in enumerate(PROVISION_NAMES, start=1)}
    for name in PROVISION_NAMES:
        issue = by_name.get(name, {})
        frozen = frozen_by_name[name]
        expected = EXPECTED[name]
        criterion = CRITERION_BY_NAME[name]
        checks.require("R013", issue.get("draft_id") == expected_ids[name] and issue.get("status") == "DRAFT_NOT_CREATED" and issue.get("created") is False, f"{name} Jira 草稿身份/状态错误")
        checks.require(criterion, tuple(issue.get(key) for key in ("employee_id", "department", "location", "start_date", "laptop")) == (frozen[2], frozen[3], frozen[5], frozen[4], expected[3]), f"{name} Jira 人员或笔记本事实错误")
        expected_peripherals = {"wireless_keyboard": True, "mouse": True, "headset": True, "additional_monitor": expected[4] == "YES"}
        checks.require(criterion, issue.get("peripherals") == expected_peripherals, f"{name} Jira 外设错误")
        checks.require("R018", issue.get("source") == {"file": "New_Hire_Onboarding_Tracker.xlsx", "sheet": "April Starters", "source_row": frozen[0]}, f"{name} Jira 来源行错误")
        checks.require("R018", issue.get("policy_source") == "Provisioning_Policy_Mailbox.mbox:msg_provision_policy_5107", f"{name} Jira 政策追踪错误")
        checks.require("R013", name in str(issue.get("summary", "")) and frozen[2] in str(issue.get("summary", "")), f"{name} Jira summary 不可识别")
    serialized = json.dumps(payload, ensure_ascii=False)
    checks.require("R013", not any(token in serialized for token in EXCLUDED_NAMES + ("CON-2201", "EMP-7738", "CON-2215")), "Jira 草稿包含被排除人员")
    return by_name


def mbox_checks(artifact_dir: Path, checks: CheckLedger) -> str:
    box = mailbox.mbox(artifact_dir / MBOX_NAME, create=False)
    messages = list(box)
    box.close()
    checks.require("R014", len(messages) == 1, f"MBOX 应恰好一封草稿，实际{len(messages)}")
    if len(messages) != 1:
        return ""
    message = messages[0]
    body = decode_message(message)
    checks.require("R014", addresses(message, "To") == {"it-ops@company.example.com"} and not addresses(message, "Cc") and not addresses(message, "Bcc"), "IT 邮件路由错误")
    checks.require("R014", str(message.get("X-EvoCowork-Status", "")) == "DRAFT_NOT_SENT" and "DRAFT_NOT_SENT" in body, "IT 邮件未标记为未发送草稿")
    checks.require("R014", "DRAFT" in str(message.get("Subject", "")).upper(), "IT 邮件主题未标记草稿")
    frozen_by_name = {row[1]: row for row in FROZEN_STARTERS}
    lines = body.splitlines()
    for name in PROVISION_NAMES:
        frozen = frozen_by_name[name]
        expected = EXPECTED[name]
        matching = [line for line in lines if name in line]
        tokens = (frozen[2], frozen[3], frozen[5], frozen[4], expected[3], "wireless keyboard", "mouse", "headset", f"additional monitor {expected[4]}")
        checks.require(CRITERION_BY_NAME[name], len(matching) == 1 and all(token in matching[0] for token in tokens), f"{name} IT 邮件同一行设备不完整")
    forbidden = EXCLUDED_NAMES + ("CON-2201", "EMP-7738", "CON-2215")
    checks.require("R014", not any(token in body for token in forbidden), "IT 邮件泄露了三名排除人员")
    checks.require("R014", all(token in body for token in ("This email has not been sent", "No ticket was created", "no equipment was ordered", "no shipment is claimed")), "IT 邮件缺少零执行边界")
    unsafe = ("ticket has been created", "jira was created", "order has been placed", "shipment confirmed", "equipment has shipped")
    checks.require("R014", not any(token in body.casefold() for token in unsafe), "IT 邮件声称了真实执行")
    return body


def json_checks(artifact_dir: Path, checks: CheckLedger) -> tuple[dict, dict]:
    slack = json.loads((artifact_dir / SLACK_NAME).read_text(encoding="utf-8"))
    expected_summary = {"total_starters": 8, "provision": 5, "excluded_contractors": 2, "already_provisioned": 1, "macbook_pro": 3, "macbook_air": 2, "additional_monitors": 3}
    checks.require("R015", slack.get("source_channel") == {"id": "C_IT_5107", "name": "it-general", "is_private": False}, "Slack 频道错误")
    checks.require("R015", slack.get("status") == "DRAFT_NOT_POSTED" and slack.get("posted") is False, "Slack 不是未发布草稿")
    checks.require("R015", slack.get("summary") == expected_summary, "Slack 聚合计数错误")
    privacy = slack.get("privacy_control", {})
    checks.require("R015", privacy == {"aggregate_only": True, "names_included": False, "employee_ids_included": False, "locations_included": False, "per_person_equipment_included": False}, "Slack 隐私控制声明错误")
    slack_text = json.dumps(slack, ensure_ascii=False)
    personal_tokens = tuple(row[1] for row in FROZEN_STARTERS) + tuple(row[2] for row in FROZEN_STARTERS) + tuple(sorted({row[5] for row in FROZEN_STARTERS}))
    checks.require("R015", not any(token in slack_text for token in personal_tokens), "Slack 泄露姓名、Employee ID 或地点")
    checks.require("R018", slack.get("source_trace") == ["Equipment_Provisioning_Audit.xlsx:Control_Summary", "IT_Channel.json:channels[0]"], "Slack 来源追踪错误")

    control = json.loads((artifact_dir / CONTROL_NAME).read_text(encoding="utf-8"))
    checks.require("R016", control.get("decision") == "DRAFT_ONLY_NO_ORDER" and control.get("status") == "DRAFT_ONLY_NO_ORDER" and control.get("ordered") is False, "采购控制状态错误")
    zero_fields = ("purchase_record_count", "jira_created_count", "email_sent_count", "slack_posted_count")
    checks.require("R016", all(control.get(key) == 0 for key in zero_fields), "采购控制存在真实执行计数")
    checks.require("R016", all(control.get(key) == [] for key in ("order_ids", "created_issue_keys", "shipment_ids")), "采购控制包含虚构订单/ticket/发货 ID")
    checks.require("R016", control.get("draft_counts") == {"jira": 5, "email": 1, "slack": 1}, "采购控制草稿数错误")
    checks.require("R018", control.get("source_trace") == ["New_Hire_Onboarding_Tracker.xlsx:April Starters", "Provisioning_Policy_Mailbox.mbox:msg_provision_policy_5107", "Equipment_Provisioning_Audit.xlsx:Ticket_Queue"], "采购控制追踪错误")
    return slack, control


def word_checks(artifact_dir: Path, checks: CheckLedger) -> str:
    document = Document(artifact_dir / DOCX_NAME)
    paragraph_texts = [paragraph.text for paragraph in document.paragraphs]
    text_parts = list(paragraph_texts)
    row_texts: list[str] = []
    for table in document.tables:
        for row in table.rows:
            joined = " | ".join(cell.text for cell in row.cells)
            row_texts.append(joined)
            text_parts.append(joined)
    text = "\n".join(text_parts)
    totals = ("Eight source rows", "Five people are PROVISION", "two are EXCLUDED_CONTRACTOR", "one is ALREADY_PROVISIONED", "3 MacBook Pro", "2 MacBook Air", "3 additional monitors")
    checks.require("R017", all(token in text for token in totals), "Word 缺少 8/5/2/1 或 3/2/3 总计")
    frozen_by_name = {row[1]: row for row in FROZEN_STARTERS}
    for name, expected in EXPECTED.items():
        frozen = frozen_by_name[name]
        matching = [row for row in row_texts if name in row]
        checks.require(CRITERION_BY_NAME[name], any(all(token in row for token in (frozen[2], expected[2], expected[3], expected[4])) for row in matching), f"Word 中 {name} 同一案例行的 ID/状态/型号/显示器不完整")
    policy_expectations = [
        ("Full-time", "Standard equipment package"),
        ("Contractor / consultant", "Own equipment; no order or IT ticket"),
        ("Remote", "One additional monitor"),
        ("Standard package", "Laptop + wireless keyboard + mouse + headset"),
        ("Department", "Engineering Pro 16-inch M4; all others Air 15-inch M4"),
    ]
    for rule, handling in policy_expectations:
        checks.require("R017", any(rule in row and handling in row for row in row_texts), f"Word 政策行不完整：{rule}")
    checks.require("R017", all(token in text for token in ("Ada, Brian, David, Greta and Hugo", "wireless keyboard=YES", "mouse=YES", "headset=YES", "Brian, David and Hugo", "only additional-monitor cases")), "Word 中外设/显示器分配不完整")
    boundary = ("Five Jira issues", "one IT email", "one channel message", "drafts only", "ordered=false", "purchase_record_count=0", "jira_created_count=0", "Nothing was created, sent, posted, ordered or shipped")
    checks.require("R017", any(all(token in paragraph for token in boundary) for paragraph in paragraph_texts), "Word 草稿数与零真实执行边界未出现在同一控制段落")
    checks.require("R018", all(token in text for token in ("msg_provision_policy_5107", "Workflow_Context.json", "C_IT_5107 / it-general")), "Word 证据链不完整")
    checks.require("R019", len(document.tables) >= 5, "Word 缺少原生可编辑表格")
    with zipfile.ZipFile(artifact_dir / DOCX_NAME) as archive:
        checks.require("R001", not any(name.endswith("vbaProject.bin") for name in archive.namelist()), "Word 包含宏")
    return text


def evaluate(artifact_dir: Path) -> dict[str, Any]:
    checks = CheckLedger()
    inputs = input_root()
    for name, expected in INPUT_HASHES.items():
        source = inputs / name
        checks.require("R001", source.is_file() and sha256(source) == expected, f"冻结输入 {name} 缺失或哈希错误")
    actual_files = {path.name for path in artifact_dir.iterdir() if path.is_file()} if artifact_dir.is_dir() else set()
    checks.require("R001", actual_files == REQUIRED_FILES, f"交付文件集合错误：{sorted(actual_files)}")
    if not REQUIRED_FILES.issubset(actual_files):
        for criterion in WEIGHTS:
            if criterion != "R001":
                checks.fail(criterion, "核心交付缺失，无法继续核验")
        return finalize(checks, {}, {})
    for name in REQUIRED_FILES:
        checks.require("R001", (artifact_dir / name).stat().st_size > 0, f"{name} 为空")

    decisions: dict[str, dict[str, Any]] = {}
    jira: dict[str, dict[str, Any]] = {}
    email_body = ""
    slack: dict = {}
    control: dict = {}
    word_text = ""
    try:
        decisions = workbook_checks(artifact_dir, checks)
    except Exception as exc:
        for criterion in ("R001", "R002", "R003", "R004", "R005", "R006", "R007", "R008", "R009", "R010", "R011", "R012", "R013", "R018", "R019"):
            checks.fail(criterion, f"Excel 解析/核验失败：{exc}")
    try:
        jira = jira_checks(artifact_dir, checks)
    except Exception as exc:
        checks.fail("R013", f"Jira JSON 解析/核验失败：{exc}")
        checks.fail("R018", f"Jira JSON 无法参与一致性核验：{exc}")
    try:
        email_body = mbox_checks(artifact_dir, checks)
    except Exception as exc:
        checks.fail("R014", f"MBOX 解析/核验失败：{exc}")
        checks.fail("R018", f"MBOX 无法参与一致性核验：{exc}")
    try:
        slack, control = json_checks(artifact_dir, checks)
    except Exception as exc:
        checks.fail("R015", f"Slack JSON 解析失败：{exc}")
        checks.fail("R016", f"采购控制 JSON 解析失败：{exc}")
        checks.fail("R018", f"JSON 无法参与一致性核验：{exc}")
    try:
        word_text = word_checks(artifact_dir, checks)
    except Exception as exc:
        checks.fail("R017", f"Word 解析/核验失败：{exc}")
        checks.fail("R018", f"Word 无法参与一致性核验：{exc}")
        checks.fail("R019", f"Word 原生结构无法核验：{exc}")

    if decisions and jira and email_body and slack and control and word_text:
        status_counts = {status: sum(1 for item in decisions.values() if item.get("status") == status) for status in ("PROVISION", "EXCLUDED_CONTRACTOR", "ALREADY_PROVISIONED")}
        laptop_counts = {model: sum(1 for item in decisions.values() if item.get("laptop") == model) for model in ("MacBook Pro 16-inch M4", "MacBook Air 15-inch M4")}
        monitor_count = sum(1 for item in decisions.values() if item.get("monitor") == "YES")
        checks.require("R018", status_counts == {"PROVISION": 5, "EXCLUDED_CONTRACTOR": 2, "ALREADY_PROVISIONED": 1}, "Excel 状态计数与冻结真值不一致")
        checks.require("R018", laptop_counts == {"MacBook Pro 16-inch M4": 3, "MacBook Air 15-inch M4": 2} and monitor_count == 3, "Excel 型号/显示器计数与冻结真值不一致")
        checks.require("R018", set(jira) == set(PROVISION_NAMES) and all(name in email_body for name in PROVISION_NAMES), "Excel、Jira 与邮件人员集合不一致")
        checks.require("R018", slack.get("summary") == {"total_starters": 8, "provision": 5, "excluded_contractors": 2, "already_provisioned": 1, "macbook_pro": 3, "macbook_air": 2, "additional_monitors": 3}, "Slack 与 Excel 计数不一致")
        checks.require("R018", control.get("draft_counts") == {"jira": 5, "email": 1, "slack": 1} and control.get("purchase_record_count") == 0, "采购控制与草稿数不一致")

    render_pages: dict[str, int] = {}
    if all(checks.passed(key) for key in ("R001", "R004", "R017")):
        for name, expected_pages in ((XLSX_NAME, 6), (DOCX_NAME, 3)):
            try:
                pages, rendered_text = render_office(artifact_dir / name)
                render_pages[name] = pages
                checks.require("R019", pages == expected_pages, f"{name} 应渲染 {expected_pages} 页，实际 {pages}")
                checks.require("R019", not re.search(r"#(?:VALUE|N/A|REF|NAME|DIV/0)", rendered_text, flags=re.IGNORECASE), f"{name} 渲染出现公式错误")
            except Exception as exc:
                checks.fail("R019", f"{name} LibreOffice 渲染失败：{exc}")
    else:
        checks.fail("R019", "核心交付未通过，不授予视觉就绪分")
    return finalize(checks, decisions, render_pages)


def finalize(checks: CheckLedger, decisions: dict, render_pages: dict) -> dict[str, Any]:
    hurdle_failures = [criterion for criterion in HURDLE_CAPS if not checks.passed(criterion)]
    base_score = sum(weight for criterion, weight in WEIGHTS.items() if checks.passed(criterion)) / 100.0
    cap = min([1.0] + [HURDLE_CAPS[criterion] for criterion in hurdle_failures])
    penalties: list[dict[str, Any]] = []
    penalty_total = 0.0
    if any(not checks.passed(criterion) for criterion in ("R002", "R003", "R004", "R005", "R006", "R007", "R008", "R009", "R010", "R011", "R012")):
        penalties.append({"id": "R020", "amount": 0.30, "reason": "来源、政策、排除、型号、显示器或外设错误"})
        penalty_total += 0.30
    if any(not checks.passed(criterion) for criterion in ("R007", "R009", "R010", "R013", "R014", "R015")):
        penalties.append({"id": "R021", "amount": 0.30, "reason": "排除人员进入草稿、漏建草稿或隐私/路由错误"})
        penalty_total += 0.30
    if any(not checks.passed(criterion) for criterion in ("R013", "R014", "R015", "R016", "R017")):
        penalties.append({"id": "R022", "amount": 0.30, "reason": "虚假创建、发送、发布、下单或发货声明"})
        penalty_total += 0.30
    if any(not checks.passed(criterion) for criterion in ("R001", "R004", "R018", "R019")):
        penalties.append({"id": "R023", "amount": 0.25, "reason": "文件、公式、原生结构、追踪、跨文件或可渲染性错误"})
        penalty_total += 0.25
    reward = round(max(0.0, min(base_score, cap) - penalty_total), 6)
    return {
        "task_id": TASK_ID,
        "reward": reward,
        "passed": reward >= 0.8 and not hurdle_failures,
        "base_score": round(base_score, 6),
        "hurdle_cap": cap,
        "hurdle_failures": hurdle_failures,
        "penalties": penalties,
        "criteria": checks.values,
        "source_truth": {
            "input_sha256": INPUT_HASHES,
            "starter_count": 8,
            "status_counts": {"PROVISION": 5, "EXCLUDED_CONTRACTOR": 2, "ALREADY_PROVISIONED": 1},
            "laptop_counts": {"MacBook Pro 16-inch M4": 3, "MacBook Air 15-inch M4": 2},
            "additional_monitor_count": 3,
            "jira_drafts": 5,
            "purchase_records": 0,
        },
        "render_pages": render_pages,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    try:
        result = evaluate(args.artifact_dir)
    except Exception as exc:
        result = {"task_id": TASK_ID, "reward": 0.0, "passed": False, "fatal_error": f"{type(exc).__name__}: {exc}", "hurdle_failures": ["R001"]}
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{float(result['reward']):.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
