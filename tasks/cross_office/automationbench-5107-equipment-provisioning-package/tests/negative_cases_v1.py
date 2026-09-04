#!/usr/bin/env python3
"""对 AutomationBench 5107 Golden 注入本题专属错误并验证 Judge 拒绝。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
import zipfile
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path
from typing import Callable

from docx import Document
from lxml import etree
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
XLSX = "Equipment_Provisioning_Audit.xlsx"
DOCX = "Provisioning_Control_Brief.docx"
JIRA = "Provisioning_Jira_Drafts.json"
MBOX = "IT_Equipment_List_Draft.mbox"
SLACK = "IT_Channel_Draft.json"
CONTROL = "Purchase_Execution_Control.json"
Mutator = Callable[[Path], None]


def load_judge():
    spec = importlib.util.spec_from_file_location("equipment_5107_judge", ROOT / "tests" / "judge.py")
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入本题 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def rewrite_zip(path: Path, transform: Callable[[str, bytes], bytes | None]) -> None:
    temporary = path.with_suffix(path.suffix + ".negative")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = transform(info.filename, source.read(info.filename))
            if data is not None:
                target.writestr(info, data)
    temporary.replace(path)


def replace_ooxml_text(path: Path, old: str, new: str) -> None:
    old_bytes, new_bytes = old.encode("utf-8"), new.encode("utf-8")
    changed = False

    def transform(name: str, data: bytes) -> bytes:
        nonlocal changed
        if name.endswith(".xml") and old_bytes in data:
            changed = True
            return data.replace(old_bytes, new_bytes)
        return data

    rewrite_zip(path, transform)
    if not changed:
        raise ValueError(f"OOXML 中找不到 {old!r}")


def sheet_xml_path(archive: zipfile.ZipFile, sheet_name: str) -> str:
    workbook = etree.fromstring(archive.read("xl/workbook.xml"))
    relationships = etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    ns = {
        "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    rel_id = workbook.xpath(f'//m:sheet[@name="{sheet_name}"]/@r:id', namespaces=ns)[0]
    target = relationships.xpath(f'//pr:Relationship[@Id="{rel_id}"]/@Target', namespaces=ns)[0]
    return target.lstrip("/") if target.startswith("/") else "xl/" + target.lstrip("/")


def edit_cell(path: Path, sheet_name: str, coordinate: str, *, value: str | int | None = None, remove_formula: bool = False) -> None:
    with zipfile.ZipFile(path, "r") as archive:
        target_sheet = sheet_xml_path(archive, sheet_name)
    changed = False
    ns_uri = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

    def transform(name: str, data: bytes) -> bytes:
        nonlocal changed
        if name != target_sheet:
            return data
        root = etree.fromstring(data)
        cells = root.xpath(f'//m:c[@r="{coordinate}"]', namespaces={"m": ns_uri})
        if not cells:
            raise ValueError(f"{sheet_name}!{coordinate} 不存在")
        cell = cells[0]
        if remove_formula:
            formula = cell.find(f"{{{ns_uri}}}f")
            if formula is None:
                raise ValueError(f"{sheet_name}!{coordinate} 没有公式")
            cell.remove(formula)
        if value is not None:
            node = cell.find(f"{{{ns_uri}}}v")
            if node is None:
                node = etree.SubElement(cell, f"{{{ns_uri}}}v")
            node.text = str(value)
            if isinstance(value, int):
                cell.attrib.pop("t", None)
            else:
                cell.set("t", "str")
        changed = True
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    rewrite_zip(path, transform)
    if not changed:
        raise ValueError("目标单元格未修改")


def mutate_json(path: Path, mutator: Callable[[dict], None]) -> None:
    payload = json.loads(path.read_text(encoding="utf-8"))
    mutator(payload)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def decode_message(message) -> str:
    payload = message.get_payload(decode=True) or b""
    return payload.decode(message.get_content_charset() or "utf-8", errors="replace")


def mutate_mbox(path: Path, mutator: Callable[[dict], None]) -> None:
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    if len(messages) != 1:
        raise ValueError("基准 MBOX 不是一封邮件")
    source = messages[0]
    state = {
        "from": str(source.get("From", "")), "to": str(source.get("To", "")),
        "cc": str(source.get("Cc", "")), "bcc": str(source.get("Bcc", "")),
        "subject": str(source.get("Subject", "")), "date": str(source.get("Date", "")),
        "message_id": str(source.get("Message-ID", "")), "status": str(source.get("X-EvoCowork-Status", "")),
        "source": str(source.get("X-EvoCowork-Source", "")), "body": decode_message(source),
    }
    mutator(state)
    message = EmailMessage(policy=SMTP)
    for key, header in (("from", "From"), ("to", "To"), ("cc", "Cc"), ("bcc", "Bcc"), ("subject", "Subject"), ("date", "Date"), ("message_id", "Message-ID"), ("status", "X-EvoCowork-Status"), ("source", "X-EvoCowork-Source")):
        if state.get(key):
            message[header] = state[key]
    message.set_content(state["body"], charset="utf-8")
    buffer = BytesIO()
    BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
    path.write_bytes(b"From hr-ops@company.example.com Wed Apr 15 09:00:00 2026\n" + buffer.getvalue().replace(b"\r\n", b"\n") + b"\n")


def body_replace(state: dict, old: str, new: str) -> None:
    if old not in state["body"]:
        raise ValueError(f"MBOX 正文中找不到 {old!r}")
    state["body"] = state["body"].replace(old, new, 1)


def drop_person_line(state: dict, name: str) -> None:
    lines = state["body"].splitlines()
    kept = [line for line in lines if name not in line]
    if len(kept) != len(lines) - 1:
        raise ValueError(f"无法唯一删除 {name} 邮件行")
    state["body"] = "\n".join(kept) + "\n"


def replace_docx_text(path: Path, old: str, new: str, *, row_token: str | None = None) -> None:
    document = Document(path)
    changed = False
    paragraphs = list(document.paragraphs) if row_token is None else []
    for table in document.tables:
        for row in table.rows:
            joined = " | ".join(cell.text for cell in row.cells)
            if row_token is not None and row_token not in joined:
                continue
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        if old in paragraph.text:
            paragraph.text = paragraph.text.replace(old, new, 1)
            changed = True
            break
    if not changed:
        raise ValueError(f"Word 中找不到 {old!r}")
    document.save(path)


def openpyxl_mutation(path: Path, mutator: Callable) -> None:
    workbook = load_workbook(path)
    mutator(workbook)
    workbook.save(path)


def build_cases() -> list[tuple[str, Mutator]]:
    cases: list[tuple[str, Mutator]] = [
        ("missing-docx", lambda root: (root / DOCX).unlink()),
        ("extra-debug-file", lambda root: (root / "debug.txt").write_text("debug\n", encoding="utf-8")),
        ("corrupt-xlsx", lambda root: (root / XLSX).write_bytes(b"not-an-xlsx")),
        ("source-name-changed", lambda root: replace_ooxml_text(root / XLSX, "Ada Petrova", "Ada Petrov")),
        ("source-id-changed", lambda root: replace_ooxml_text(root / XLSX, "EMP-7715", "EMP-7716")),
        ("source-department-changed", lambda root: replace_ooxml_text(root / XLSX, "Support", "Engineering")),
        ("source-location-changed", lambda root: replace_ooxml_text(root / XLSX, "Remote - Madrid", "Office - Madrid")),
        ("source-employment-type-changed", lambda root: replace_ooxml_text(root / XLSX, "Independent Contractor", "Full-time")),
        ("source-shipped-note-changed", lambda root: replace_ooxml_text(root / XLSX, "Equipment already shipped from previous office transfer", "Equipment requested")),
        ("policy-full-time-changed", lambda root: replace_ooxml_text(root / XLSX, "Receive the standard equipment package", "Receive no equipment")),
        ("policy-contractor-changed", lambda root: replace_ooxml_text(root / XLSX, "Provide their own equipment; no order or IT ticket", "Create an IT ticket")),
        ("policy-remote-changed", lambda root: replace_ooxml_text(root / XLSX, "Add one monitor", "No monitor")),
        ("policy-model-changed", lambda root: replace_ooxml_text(root / XLSX, "Engineering: MacBook Pro 16-inch M4; all others: MacBook Air 15-inch M4", "Everyone receives MacBook Air")),
        ("evidence-time-changed", lambda root: replace_ooxml_text(root / XLSX, "2026-04-15T09:00:00Z", "2026-04-16T09:00:00Z")),
        ("evidence-channel-changed", lambda root: replace_ooxml_text(root / XLSX, "C_IT_5107", "C_IT_WRONG")),
        ("formula-note-staticized", lambda root: edit_cell(root / XLSX, "Provisioning_Decisions", "H2", remove_formula=True)),
        ("formula-contractor-staticized", lambda root: edit_cell(root / XLSX, "Provisioning_Decisions", "I4", remove_formula=True)),
        ("formula-status-staticized", lambda root: edit_cell(root / XLSX, "Provisioning_Decisions", "K6", remove_formula=True)),
        ("formula-laptop-staticized", lambda root: edit_cell(root / XLSX, "Provisioning_Decisions", "L5", remove_formula=True)),
        ("formula-monitor-staticized", lambda root: edit_cell(root / XLSX, "Provisioning_Decisions", "M9", remove_formula=True)),
        ("ada-model-wrong", lambda root: edit_cell(root / XLSX, "Provisioning_Decisions", "L2", value="MacBook Air 15-inch M4")),
        ("brian-monitor-missing", lambda root: edit_cell(root / XLSX, "Provisioning_Decisions", "M3", value="NO")),
        ("carmen-provisioned", lambda root: edit_cell(root / XLSX, "Provisioning_Decisions", "K4", value="PROVISION")),
        ("david-model-wrong", lambda root: edit_cell(root / XLSX, "Provisioning_Decisions", "L5", value="MacBook Air 15-inch M4")),
        ("elena-duplicate-provision", lambda root: edit_cell(root / XLSX, "Provisioning_Decisions", "K6", value="PROVISION")),
        ("femi-provisioned", lambda root: edit_cell(root / XLSX, "Provisioning_Decisions", "K7", value="PROVISION")),
        ("greta-headset-missing", lambda root: edit_cell(root / XLSX, "Provisioning_Decisions", "P8", value="NO")),
        ("hugo-monitor-missing", lambda root: edit_cell(root / XLSX, "Provisioning_Decisions", "M9", value="NO")),
        ("summary-provision-wrong", lambda root: edit_cell(root / XLSX, "Control_Summary", "B3", value=4)),
        ("summary-orders-nonzero", lambda root: edit_cell(root / XLSX, "Control_Summary", "B11", value=1)),
        ("queue-wrong-person", lambda root: edit_cell(root / XLSX, "Ticket_Queue", "B3", value="Carmen Delgado")),
        ("queue-wrong-laptop", lambda root: edit_cell(root / XLSX, "Ticket_Queue", "G4", value="MacBook Air 15-inch M4")),
        ("queue-wrong-source-row", lambda root: edit_cell(root / XLSX, "Ticket_Queue", "M6", value="New_Hire_Onboarding_Tracker.xlsx:April Starters row 8")),
        ("native-chart-removed", lambda root: openpyxl_mutation(root / XLSX, lambda wb: setattr(wb["Control_Summary"], "_charts", []))),
        ("native-table-removed", lambda root: openpyxl_mutation(root / XLSX, lambda wb: wb["Source_Starters"].tables.clear())),
        ("conditional-formatting-removed", lambda root: openpyxl_mutation(root / XLSX, lambda wb: setattr(wb["Provisioning_Decisions"], "conditional_formatting", type(wb["Provisioning_Decisions"].conditional_formatting)()))),
    ]

    cases.extend([
        ("jira-top-level-created", lambda root: mutate_json(root / JIRA, lambda p: p.update({"created": True, "status": "CREATED"}))),
        ("jira-count-wrong", lambda root: mutate_json(root / JIRA, lambda p: p.update({"issue_count": 4}))),
        ("jira-fake-created-key", lambda root: mutate_json(root / JIRA, lambda p: p.update({"created_issue_keys": ["IT-999"]}))),
        ("jira-missing-hugo", lambda root: mutate_json(root / JIRA, lambda p: p["issues"].pop())),
        ("jira-add-carmen", lambda root: mutate_json(root / JIRA, lambda p: p["issues"].append({**p["issues"][0], "person": "Carmen Delgado", "employee_id": "CON-2201"}))),
        ("jira-ada-model-wrong", lambda root: mutate_json(root / JIRA, lambda p: p["issues"][0].update({"laptop": "MacBook Air 15-inch M4"}))),
        ("jira-brian-monitor-missing", lambda root: mutate_json(root / JIRA, lambda p: p["issues"][1]["peripherals"].update({"additional_monitor": False}))),
        ("jira-david-department-wrong", lambda root: mutate_json(root / JIRA, lambda p: p["issues"][2].update({"department": "Sales"}))),
        ("jira-greta-keyboard-missing", lambda root: mutate_json(root / JIRA, lambda p: p["issues"][3]["peripherals"].update({"wireless_keyboard": False}))),
        ("jira-hugo-source-row-wrong", lambda root: mutate_json(root / JIRA, lambda p: p["issues"][4]["source"].update({"source_row": 8}))),
        ("jira-issue-created", lambda root: mutate_json(root / JIRA, lambda p: p["issues"][0].update({"created": True, "status": "CREATED"}))),
        ("jira-policy-trace-missing", lambda root: mutate_json(root / JIRA, lambda p: p["issues"][2].update({"policy_source": ""}))),
        ("jira-start-date-wrong", lambda root: mutate_json(root / JIRA, lambda p: p["issues"][3].update({"start_date": "2026-04-22"}))),
        ("jira-summary-excluded-leak", lambda root: mutate_json(root / JIRA, lambda p: p["issues"][0].update({"summary": p["issues"][0]["summary"] + " Carmen Delgado"}))),
    ])

    cases.extend([
        ("mbox-duplicate-message", lambda root: (root / MBOX).write_bytes((root / MBOX).read_bytes() * 2)),
        ("mbox-wrong-recipient", lambda root: mutate_mbox(root / MBOX, lambda s: s.update({"to": "hr@company.example.com"}))),
        ("mbox-add-cc", lambda root: mutate_mbox(root / MBOX, lambda s: s.update({"cc": "manager@company.example.com"}))),
        ("mbox-false-sent", lambda root: mutate_mbox(root / MBOX, lambda s: s.update({"status": "SENT"}))),
        ("mbox-subject-not-draft", lambda root: mutate_mbox(root / MBOX, lambda s: s.update({"subject": "April 21 starter equipment list"}))),
        ("mbox-drop-ada", lambda root: mutate_mbox(root / MBOX, lambda s: drop_person_line(s, "Ada Petrova"))),
        ("mbox-drop-brian", lambda root: mutate_mbox(root / MBOX, lambda s: drop_person_line(s, "Brian Osei"))),
        ("mbox-drop-david", lambda root: mutate_mbox(root / MBOX, lambda s: drop_person_line(s, "David Nguyen"))),
        ("mbox-drop-greta", lambda root: mutate_mbox(root / MBOX, lambda s: drop_person_line(s, "Greta Holm"))),
        ("mbox-drop-hugo", lambda root: mutate_mbox(root / MBOX, lambda s: drop_person_line(s, "Hugo Fernandez"))),
        ("mbox-brian-model-wrong", lambda root: mutate_mbox(root / MBOX, lambda s: body_replace(s, "Brian Osei | EMP-7715 | Sales | Remote - Chicago | start 2026-04-21 | MacBook Air 15-inch M4", "Brian Osei | EMP-7715 | Sales | Remote - Chicago | start 2026-04-21 | MacBook Pro 16-inch M4"))),
        ("mbox-ada-monitor-wrong", lambda root: mutate_mbox(root / MBOX, lambda s: body_replace(s, "Ada Petrova | EMP-7701 | Engineering | Office - Austin | start 2026-04-21 | MacBook Pro 16-inch M4 | wireless keyboard | mouse | headset | additional monitor NO", "Ada Petrova | EMP-7701 | Engineering | Office - Austin | start 2026-04-21 | MacBook Pro 16-inch M4 | wireless keyboard | mouse | headset | additional monitor YES"))),
        ("mbox-add-carmen", lambda root: mutate_mbox(root / MBOX, lambda s: s.update({"body": s["body"] + "\nCarmen Delgado | CON-2201\n"}))),
        ("mbox-claim-created", lambda root: mutate_mbox(root / MBOX, lambda s: body_replace(s, "No ticket was created", "A ticket has been created"))),
    ])

    cases.extend([
        ("slack-false-posted", lambda root: mutate_json(root / SLACK, lambda p: p.update({"posted": True, "status": "POSTED"}))),
        ("slack-wrong-channel", lambda root: mutate_json(root / SLACK, lambda p: p["source_channel"].update({"id": "C_WRONG"}))),
        ("slack-provision-count-wrong", lambda root: mutate_json(root / SLACK, lambda p: p["summary"].update({"provision": 6}))),
        ("slack-model-count-wrong", lambda root: mutate_json(root / SLACK, lambda p: p["summary"].update({"macbook_pro": 4}))),
        ("slack-monitor-count-wrong", lambda root: mutate_json(root / SLACK, lambda p: p["summary"].update({"additional_monitors": 4}))),
        ("slack-leak-name", lambda root: mutate_json(root / SLACK, lambda p: p["draft_message"].update({"text": p["draft_message"]["text"] + " Ada Petrova"}))),
        ("slack-leak-id", lambda root: mutate_json(root / SLACK, lambda p: p["draft_message"].update({"text": p["draft_message"]["text"] + " EMP-7701"}))),
        ("slack-leak-location", lambda root: mutate_json(root / SLACK, lambda p: p["draft_message"].update({"text": p["draft_message"]["text"] + " Remote - Chicago"}))),
        ("slack-privacy-flag-wrong", lambda root: mutate_json(root / SLACK, lambda p: p["privacy_control"].update({"names_included": True}))),
        ("slack-source-trace-missing", lambda root: mutate_json(root / SLACK, lambda p: p.update({"source_trace": []}))),
    ])

    cases.extend([
        ("control-false-ordered", lambda root: mutate_json(root / CONTROL, lambda p: p.update({"ordered": True, "status": "ORDERED"}))),
        ("control-purchase-count-nonzero", lambda root: mutate_json(root / CONTROL, lambda p: p.update({"purchase_record_count": 1}))),
        ("control-jira-created-nonzero", lambda root: mutate_json(root / CONTROL, lambda p: p.update({"jira_created_count": 1}))),
        ("control-email-sent-nonzero", lambda root: mutate_json(root / CONTROL, lambda p: p.update({"email_sent_count": 1}))),
        ("control-slack-posted-nonzero", lambda root: mutate_json(root / CONTROL, lambda p: p.update({"slack_posted_count": 1}))),
        ("control-fake-order-id", lambda root: mutate_json(root / CONTROL, lambda p: p.update({"order_ids": ["PO-FAKE"]}))),
        ("control-fake-ticket-key", lambda root: mutate_json(root / CONTROL, lambda p: p.update({"created_issue_keys": ["IT-123"]}))),
        ("control-fake-shipment-id", lambda root: mutate_json(root / CONTROL, lambda p: p.update({"shipment_ids": ["SHIP-1"]}))),
        ("control-draft-count-wrong", lambda root: mutate_json(root / CONTROL, lambda p: p["draft_counts"].update({"jira": 4}))),
        ("control-source-trace-missing", lambda root: mutate_json(root / CONTROL, lambda p: p.update({"source_trace": []}))),
    ])

    cases.extend([
        ("word-total-wrong", lambda root: replace_docx_text(root / DOCX, "Eight source rows", "Seven source rows")),
        ("word-ada-model-wrong", lambda root: replace_docx_text(root / DOCX, "MacBook Pro 16-inch M4", "MacBook Air 15-inch M4", row_token="Ada Petrova")),
        ("word-brian-monitor-wrong", lambda root: replace_docx_text(root / DOCX, "YES", "NO", row_token="Brian Osei")),
        ("word-carmen-status-wrong", lambda root: replace_docx_text(root / DOCX, "EXCLUDED_CONTRACTOR", "PROVISION", row_token="Carmen Delgado")),
        ("word-david-model-wrong", lambda root: replace_docx_text(root / DOCX, "MacBook Pro 16-inch M4", "MacBook Air 15-inch M4", row_token="David Nguyen")),
        ("word-elena-status-wrong", lambda root: replace_docx_text(root / DOCX, "ALREADY_PROVISIONED", "PROVISION", row_token="Elena Rossi")),
        ("word-femi-status-wrong", lambda root: replace_docx_text(root / DOCX, "EXCLUDED_CONTRACTOR", "PROVISION", row_token="Femi Adeyemi")),
        ("word-greta-model-wrong", lambda root: replace_docx_text(root / DOCX, "MacBook Air 15-inch M4", "MacBook Pro 16-inch M4", row_token="Greta Holm")),
        ("word-hugo-monitor-wrong", lambda root: replace_docx_text(root / DOCX, "YES", "NO", row_token="Hugo Fernandez")),
        ("word-contractor-policy-wrong", lambda root: replace_docx_text(root / DOCX, "Own equipment; no order or IT ticket", "Create an IT ticket", row_token="Contractor / consultant")),
        ("word-remote-policy-wrong", lambda root: replace_docx_text(root / DOCX, "One additional monitor", "No monitor", row_token="Remote")),
        ("word-standard-package-wrong", lambda root: replace_docx_text(root / DOCX, "Laptop + wireless keyboard + mouse + headset", "Laptop only", row_token="Standard package")),
        ("word-department-policy-wrong", lambda root: replace_docx_text(root / DOCX, "Engineering Pro 16-inch M4; all others Air 15-inch M4", "Everyone receives Air", row_token="Department")),
        ("word-false-order", lambda root: replace_docx_text(root / DOCX, "ordered=false", "ordered=true")),
        ("word-evidence-trace-missing", lambda root: replace_docx_text(root / DOCX, "msg_provision_policy_5107", "missing-policy")),
    ])
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    judge = load_judge()
    # Golden 的真实 LibreOffice 渲染已单独逐页验收；负例只验证语义与结构拒错。
    judge.render_office = lambda path: (6 if path.suffix.casefold() == ".xlsx" else 3, "")
    results: list[dict] = []
    for name, mutator in build_cases():
        with tempfile.TemporaryDirectory(prefix=f"negative-5107-{name}-") as temporary:
            output = Path(temporary) / "output"
            shutil.copytree(GOLDEN, output)
            mutator(output)
            result = judge.evaluate(output)
            results.append({
                "case": name,
                "reward": result.get("reward", 0.0),
                "passed": result.get("passed", False),
                "hurdle_failures": result.get("hurdle_failures", []),
                "penalties": result.get("penalties", []),
            })
    report = {
        "total": len(results),
        "rejected": sum(not item["passed"] for item in results),
        "maximum_reward": max(item["reward"] for item in results),
        "accepted_cases": [item["case"] for item in results if item["passed"]],
        "cases": results,
    }
    if args.report:
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if report["rejected"] != report["total"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
