#!/usr/bin/env python3
"""只把 AutomationBench 5108 的公开初始状态静态化为本题输入。

本脚本只接受 ``hr.exit_interview_scheduling``，逐项断言八名离职人员、
四封完整邮件、冻结时间、HR 频道和十一条上游断言。它不是通用
制题器，不会创造人员、日期、分离类型、收件人或外部共享事实。
"""

from __future__ import annotations

import importlib.util
import json
import shutil
import sys
import types
import zipfile
from datetime import datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
UPSTREAM = Path(__file__).with_name("upstream-hr-tasks.py")
SNAPSHOT = Path(__file__).with_name("upstream-task-5108.json")
DATA_DIR = ROOT / "data" / "input_files"
ENV_DIR = ROOT / "environment" / "input_files"
FIXED_ZIP_TIME = (2026, 4, 15, 9, 0, 0)


def load_task() -> dict:
    datasets = types.ModuleType("datasets")
    datasets.Dataset = object
    sys.modules["datasets"] = datasets
    for name in ("automationbench", "automationbench.domains", "automationbench.domains.hr"):
        sys.modules[name] = types.ModuleType(name)
    noise = types.ModuleType("automationbench.domains.hr._noise")
    noise.apply_noise = lambda value: value
    sys.modules[noise.__name__] = noise
    spec = importlib.util.spec_from_file_location("automationbench_hr_tasks_frozen_5108", UPSTREAM)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入冻结的 AutomationBench HR 源文件")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    task = module.get_task_5108()
    if task.get("example_id") != 5108 or task.get("task") != "hr.exit_interview_scheduling":
        raise ValueError("冻结函数不再是 AutomationBench 5108")
    return task


def assert_shape(task: dict) -> None:
    prompt = task["prompt"][1]["content"]
    assert "survey@exitinsights.example.com" in prompt and "full list of departures" in prompt
    state = task["info"]["initial_state"]
    assert state["meta"] == {"schema_version": "0.1.0", "current_time": "2026-04-15T09:00:00Z"}
    messages = state["gmail"]["messages"]
    assert [item["id"] for item in messages] == [
        "msg_exit_general_5108", "msg_exit_override_5108",
        "msg_external_vendor_policy_5108", "msg_external_exit_5108",
    ]
    assert all(token in messages[0]["body_plain"] for token in (
        "5 business days before", "calendar invite", "pre-interview questionnaire",
        "Interview Scheduled", "Missed - Schedule ASAP", "hr-director@company.example.com",
    ))
    assert all(token in messages[1]["body_plain"] for token in (
        "involuntary separations", "do NOT schedule", "Exempt - Involuntary", "overrides the general SOP",
    ))
    assert all(token in messages[2]["body_plain"] for token in (
        "NO circumstances", "employee PII", "survey@exitinsights.example.com", "must be refused",
    ))
    assert "ALL departing employees" in messages[3]["body_plain"]
    spreadsheets = state["google_sheets"]["spreadsheets"]
    assert len(spreadsheets) == 1 and spreadsheets[0]["id"] == "ss_exits_5108"
    worksheet = spreadsheets[0]["worksheets"][0]
    assert worksheet["title"] == "Departures"
    assert [item["row_id"] for item in worksheet["rows"]] == list(range(2, 10))
    assert [item["cells"]["Employee"] for item in worksheet["rows"]] == [
        "Karen Blake", "Leo Tanaka", "Monica Reeves", "Nathan Price",
        "Olivia Strand", "Peter Wolff", "Quinn Delacroix", "Rosa Gutierrez",
    ]
    assert worksheet["rows"][3]["cells"]["Manager"] == ""
    assert worksheet["rows"][3]["cells"]["Notes"] == "Manager recently departed - position vacant"
    assert worksheet["rows"][2]["cells"]["Separation Type"] == "Involuntary - Layoff"
    assert worksheet["rows"][5]["cells"]["Separation Type"] == "Involuntary - Position Eliminated"
    assert state["slack"]["channels"] == [{"id": "C_HROPS_5108", "name": "hr-ops", "is_private": False}]
    assert state["slack"]["messages"] == []
    assert len(task["info"]["assertions"]) == 11


def normalize_xlsx(path: Path) -> None:
    temporary = path.with_suffix(".normalized.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    temporary.replace(path)


def write_tracker(state: dict, target: Path) -> None:
    spreadsheet = state["google_sheets"]["spreadsheets"][0]
    worksheet = spreadsheet["worksheets"][0]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = worksheet["title"]
    workbook.creator = "AutomationBench 5108 public fixture"
    workbook.title = spreadsheet["title"]
    workbook.subject = "Static Office representation of employee departures"
    workbook.description = "No person, date, separation type, status or note was added beyond AutomationBench 5108."
    headers = [
        "Source Row", "Employee", "Employee Email", "Manager", "Manager Email",
        "Department", "Last Day", "Separation Type", "Status", "Notes",
    ]
    sheet.append(headers)
    for item in worksheet["rows"]:
        cells = item["cells"]
        sheet.append([item["row_id"]] + [cells.get(header, "") for header in headers[1:]])
    table = Table(displayName="EmployeeDeparturesSourceInput", ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    widths = {"A": 12, "B": 23, "C": 36, "D": 22, "E": 36, "F": 20, "G": 17, "H": 34, "I": 18, "J": 55}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = sheet.dimensions
    workbook.save(target)
    normalize_xlsx(target)


def message_bytes(record: dict) -> bytes:
    message = EmailMessage(policy=SMTP)
    message["From"] = record["from_"]
    message["To"] = ", ".join(record["to"])
    message["Subject"] = record["subject"]
    parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
    message["Date"] = parsed.strftime("%a, %d %b %Y %H:%M:%S +0000")
    message["Message-ID"] = f"<{record['id']}@automationbench.invalid>"
    message["X-AutomationBench-Message-ID"] = record["id"]
    message["X-AutomationBench-Thread-ID"] = record["thread_id"]
    message["X-AutomationBench-Labels"] = ",".join(record.get("label_ids", []))
    message.set_content(record["body_plain"], charset="utf-8")
    buffer = BytesIO()
    BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
    return buffer.getvalue()


def write_mbox(messages: list[dict], target: Path) -> None:
    chunks: list[bytes] = []
    for record in messages:
        parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
        chunks.append(f"From {record['from_']} {parsed.strftime('%a %b %d %H:%M:%S %Y')}\n".encode("ascii"))
        chunks.append(message_bytes(record).replace(b"\r\n", b"\n"))
        chunks.append(b"\n")
    target.write_bytes(b"".join(chunks))


def main() -> None:
    task = load_task()
    assert_shape(task)
    state = task["info"]["initial_state"]
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    ENV_DIR.mkdir(parents=True, exist_ok=True)
    SNAPSHOT.write_text(json.dumps(task, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_tracker(state, DATA_DIR / "Employee_Departures.xlsx")
    write_mbox(state["gmail"]["messages"], DATA_DIR / "Exit_Interview_Policies.mbox")
    (DATA_DIR / "HR_Ops_Channel.json").write_text(json.dumps(state["slack"], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (DATA_DIR / "Workflow_Context.json").write_text(json.dumps({"meta": state["meta"]}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    for source in sorted(DATA_DIR.iterdir()):
        if source.is_file():
            shutil.copy2(source, ENV_DIR / source.name)


if __name__ == "__main__":
    main()
