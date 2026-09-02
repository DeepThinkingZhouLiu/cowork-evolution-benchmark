#!/usr/bin/env python3
"""AutomationBench 5110 组织重组控制包的逐题参考解法。"""

from __future__ import annotations

import argparse
import hashlib
import json
import mailbox
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt as PptPt


INPUT_HASHES = {
    "Employee_Directory.xlsx": "42c477c7905f4dfd072be5c6cb7a20bc127cc275c5aeff37553069cb07b80912",
    "Approved_Org_Changes.mbox": "f463fc846844feea9354791ee4c9907dc668fe5d9ba0abb6707ef2b0d9ff4e12",
    "HR_Ops_Channel.json": "6c39fc7813952d47266716507949c4b7a9adb3c9cf3ee398722595adc36371cd",
    "Workflow_Context.json": "76c4a0657bc27f2fd830b10613ec7ceae3adcdc5798aeee0f822d27bf69bbc99",
}
XLSX_NAME = "Org_Restructure_Audit.xlsx"
DOCX_NAME = "Org_Restructure_Control_Brief.docx"
PPTX_NAME = "Org_Restructure_Chart.pptx"
MBOX_NAME = "Org_Change_Notification_Drafts.mbox"
UPDATES_NAME = "Directory_Update_Drafts.json"
SLACK_NAME = "HR_Ops_Reorg_Draft.json"
SCOPE_NAME = "Scope_Control.json"
FIXED_ISO_TIME = "2026-04-15T09:00:00Z"
EFFECTIVE_DATE = "2026-04-14"
FIXED_ZIP_TIME = (2026, 4, 15, 9, 0, 0)

NAVY = "17365D"; BLUE = "5B9BD5"; GREEN = "E2F0D9"; AMBER = "FFF2CC"; RED = "F4CCCC"; GRAY = "E7E6E6"; WHITE = "FFFFFF"
THIN = Side(style="thin", color="B7C9D6")
PPT_NAVY = PptRGBColor(23, 54, 93); PPT_BLUE = PptRGBColor(91, 155, 213); PPT_LIGHT_BLUE = PptRGBColor(221, 235, 247)
PPT_GREEN = PptRGBColor(226, 240, 217); PPT_AMBER = PptRGBColor(255, 242, 204); PPT_GRAY = PptRGBColor(242, 242, 242); PPT_WHITE = PptRGBColor(255, 255, 255)


@dataclass(frozen=True)
class Employee:
    source_row: int
    employee: str
    employee_email: str
    department: str
    manager: str
    manager_email: str


@dataclass(frozen=True)
class Change:
    order: int
    employee: str
    old_manager: str
    new_manager: str
    new_manager_email: str


@dataclass(frozen=True)
class Decision:
    source: Employee
    affected: str
    approved_old_manager: str
    old_manager_matches: str
    approved_new_manager: str
    new_manager_email: str
    updated_manager: str
    updated_manager_email: str
    result: str
    directory_update: str
    notification: str
    slack_summary: str


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""): digest.update(chunk)
    return digest.hexdigest()


def decode_message(message) -> str:
    if message.is_multipart():
        parts: list[str] = []
        for part in message.walk():
            if part.get_content_type() == "text/plain":
                payload = part.get_payload(decode=True) or b""; parts.append(payload.decode(part.get_content_charset() or "utf-8", errors="replace"))
        return "\n".join(parts).strip()
    payload = message.get_payload(decode=True) or b""
    return payload.decode(message.get_content_charset() or "utf-8", errors="replace").strip()


def read_inputs(input_dir: Path) -> tuple[list[Employee], dict, dict, str, list[Change]]:
    for name, expected in INPUT_HASHES.items():
        source = input_dir / name
        if not source.is_file() or sha256(source) != expected: raise ValueError(f"输入 {name} 缺失或 SHA256 与冻结版本不一致")
    context = json.loads((input_dir / "Workflow_Context.json").read_text(encoding="utf-8"))
    if context != {"meta": {"schema_version": "0.1.0", "current_time": FIXED_ISO_TIME}}: raise ValueError("冻结运行时间已变化")
    channel = json.loads((input_dir / "HR_Ops_Channel.json").read_text(encoding="utf-8"))
    if channel != {"channels": [{"id": "C_HROPS_5110", "name": "hr-ops", "is_private": False}], "messages": []}: raise ValueError("HR Ops 频道快照已变化")

    workbook = load_workbook(input_dir / "Employee_Directory.xlsx", read_only=True, data_only=True)
    if workbook.sheetnames != ["Employees"]: raise ValueError("员工目录结构已变化")
    employees: list[Employee] = []
    for row in workbook["Employees"].iter_rows(min_row=2, values_only=True):
        if row[0] is None: continue
        employees.append(Employee(int(row[0]), str(row[1]), str(row[2]), str(row[3]), str(row[4]), str(row[5])))
    workbook.close()
    frozen = [
        (2, "Aisha Jordan", "aisha.jordan@company.example.com", "Engineering", "Lisa Wang", "lisa.wang@company.example.com"),
        (3, "Bo Lindqvist", "bo.lindqvist@company.example.com", "Engineering", "Marcus Chen", "marcus.chen@company.example.com"),
        (4, "Chiara Russo", "chiara.russo@company.example.com", "Engineering", "Derek Huang", "derek.huang@company.example.com"),
        (5, "Dante Obi", "dante.obi@company.example.com", "Product", "Lisa Wang", "lisa.wang@company.example.com"),
        (6, "Elena Vasquez", "elena.vasquez@company.example.com", "Engineering", "Derek Huang", "derek.huang@company.example.com"),
        (7, "Felix Braun", "felix.braun@company.example.com", "Engineering", "Marcus Chen", "marcus.chen@company.example.com"),
        (8, "Grace Kim", "grace.kim@company.example.com", "Product", "Marcus Chen-Williams", "marcus.chen-williams@company.example.com"),
        (9, "Hugo Fernandez", "hugo.fernandez@company.example.com", "Engineering", "Lisa Wang", "lisa.wang@company.example.com"),
    ]
    if [(e.source_row, e.employee, e.employee_email, e.department, e.manager, e.manager_email) for e in employees] != frozen: raise ValueError("八条冻结员工目录已变化")

    box = mailbox.mbox(input_dir / "Approved_Org_Changes.mbox", create=False); messages = list(box); box.close()
    if len(messages) != 1: raise ValueError("批准邮箱必须恰好包含一封邮件")
    message = messages[0]
    approved = {"message_id": str(message.get("X-AutomationBench-Message-ID", "")), "thread_id": str(message.get("X-AutomationBench-Thread-ID", "")), "from": str(message.get("From", "")), "to": str(message.get("To", "")), "subject": str(message.get("Subject", "")), "date": str(message.get("Date", "")), "body": decode_message(message)}
    if approved["message_id"] != "msg_reorg_5110" or approved["from"] != "vp-people@company.example.com" or "Effective April 14" not in approved["subject"]: raise ValueError("批准邮件身份已变化")
    manager_emails: dict[str, set[str]] = {}
    for employee in employees: manager_emails.setdefault(employee.manager, set()).add(employee.manager_email)
    if manager_emails != {"Lisa Wang": {"lisa.wang@company.example.com"}, "Marcus Chen": {"marcus.chen@company.example.com"}, "Derek Huang": {"derek.huang@company.example.com"}, "Marcus Chen-Williams": {"marcus.chen-williams@company.example.com"}}: raise ValueError("经理邮箱证据不唯一")
    changes = [
        Change(1, "Aisha Jordan", "Lisa Wang", "Marcus Chen", "marcus.chen@company.example.com"),
        Change(2, "Bo Lindqvist", "Marcus Chen", "Derek Huang", "derek.huang@company.example.com"),
        Change(3, "Chiara Russo", "Derek Huang", "Lisa Wang", "lisa.wang@company.example.com"),
        Change(4, "Dante Obi", "Lisa Wang", "Marcus Chen-Williams", "marcus.chen-williams@company.example.com"),
    ]
    for change in changes:
        line = f"{change.order}. {change.employee} -> moves from {change.old_manager} to {change.new_manager}"
        if line not in approved["body"]: raise ValueError(f"批准邮件缺少精确改线：{change.employee}")
    if "Marcus Chen (Engineering) and Marcus Chen-Williams (Product) are different managers" not in approved["body"]: raise ValueError("相似经理消歧证据缺失")
    return employees, approved, channel, context["meta"]["current_time"], changes


def derive_decisions(employees: list[Employee], changes: list[Change]) -> list[Decision]:
    change_by_name = {change.employee: change for change in changes}
    decisions: list[Decision] = []
    for employee in employees:
        change = change_by_name.get(employee.employee)
        if change:
            if employee.manager != change.old_manager: raise ValueError(f"{employee.employee} 的批准旧经理与目录冲突")
            decisions.append(Decision(employee, "YES", change.old_manager, "YES", change.new_manager, change.new_manager_email, change.new_manager, change.new_manager_email, "APPROVED_CHANGE", "YES", "YES", "YES"))
        else:
            decisions.append(Decision(employee, "NO", "", "N/A", "", "", employee.manager, employee.manager_email, "UNCHANGED_NOT_IN_APPROVED_LIST", "NO", "NO", "NO"))
    expected = {
        "Aisha Jordan": ("YES", "Lisa Wang", "Marcus Chen", "marcus.chen@company.example.com", "APPROVED_CHANGE"),
        "Bo Lindqvist": ("YES", "Marcus Chen", "Derek Huang", "derek.huang@company.example.com", "APPROVED_CHANGE"),
        "Chiara Russo": ("YES", "Derek Huang", "Lisa Wang", "lisa.wang@company.example.com", "APPROVED_CHANGE"),
        "Dante Obi": ("YES", "Lisa Wang", "Marcus Chen-Williams", "marcus.chen-williams@company.example.com", "APPROVED_CHANGE"),
        "Elena Vasquez": ("NO", "", "Derek Huang", "derek.huang@company.example.com", "UNCHANGED_NOT_IN_APPROVED_LIST"),
        "Felix Braun": ("NO", "", "Marcus Chen", "marcus.chen@company.example.com", "UNCHANGED_NOT_IN_APPROVED_LIST"),
        "Grace Kim": ("NO", "", "Marcus Chen-Williams", "marcus.chen-williams@company.example.com", "UNCHANGED_NOT_IN_APPROVED_LIST"),
        "Hugo Fernandez": ("NO", "", "Lisa Wang", "lisa.wang@company.example.com", "UNCHANGED_NOT_IN_APPROVED_LIST"),
    }
    actual = {d.source.employee: (d.affected, d.approved_old_manager, d.updated_manager, d.updated_manager_email, d.result) for d in decisions}
    if actual != expected: raise ValueError("八人组织改线真值不一致")
    return decisions


def add_table(sheet, name: str) -> None:
    table = Table(displayName=name, ref=sheet.dimensions); table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True); sheet.add_table(table)


def style_sheet(sheet, widths: dict[str, float], footer: str, *, a3: bool = False) -> None:
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions; sheet.sheet_view.showGridLines = False
    for column, width in widths.items(): sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY); cell.font = Font(name="Aptos", bold=True, color=WHITE, size=8.3); cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = Border(bottom=THIN)
    sheet.row_dimensions[1].height = 31; grid = Side(style="hair", color="D9E2F3")
    for row in sheet.iter_rows(min_row=2):
        for cell in row: cell.font = Font(name="Aptos", size=8.1); cell.alignment = Alignment(vertical="top", wrap_text=True); cell.border = Border(left=grid, right=grid, bottom=grid)
    sheet.page_setup.orientation = "landscape"; sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if a3 else sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 1; sheet.sheet_properties.pageSetUpPr.fitToPage = True; sheet.print_options.horizontalCentered = True
    sheet.oddFooter.center.text = footer; sheet.print_area = sheet.dimensions


def normalize_ooxml(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="reorg-5110-ooxml-") as temporary:
        target = Path(temporary) / path.name; axis_ids: dict[str, str] = {}
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as destination:
            for name in sorted(source.namelist()):
                data = source.read(name)
                if name == "docProps/core.xml":
                    try:
                        root = etree.fromstring(data); ns = {"dcterms": "http://purl.org/dc/terms/", "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"}
                        for tag in ("created", "modified"):
                            for node in root.xpath(f"//dcterms:{tag}", namespaces=ns): node.text = "2026-04-15T09:00:00Z"
                        for tag in ("lastModifiedBy", "revision"):
                            for node in root.xpath(f"//cp:{tag}", namespaces=ns): node.text = "EvoCowork" if tag == "lastModifiedBy" else "1"
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except etree.XMLSyntaxError: pass
                if name.startswith("xl/charts/chart") and name.endswith(".xml"):
                    try:
                        root = etree.fromstring(data); ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
                        for node in root.xpath("//c:axId", namespaces=ns):
                            old = node.get("val")
                            if old and old not in axis_ids: axis_ids[old] = str(51_100_001 + len(axis_ids))
                        for node in root.xpath("//c:axId | //c:crossAx", namespaces=ns):
                            old = node.get("val")
                            if old in axis_ids: node.set("val", axis_ids[old])
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except etree.XMLSyntaxError: pass
                info = zipfile.ZipInfo(name, FIXED_ZIP_TIME); info.compress_type = zipfile.ZIP_DEFLATED; info.external_attr = 0o600 << 16; destination.writestr(info, data)
        shutil.copy2(target, path)


def recalculate_xlsx(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="reorg-5110-recalc-") as temporary:
        root = Path(temporary); source_dir, output_dir, profile_dir = root / "source", root / "output", root / "profile"; source_dir.mkdir(); output_dir.mkdir()
        source = source_dir / path.name; shutil.copy2(path, source)
        process = subprocess.run(["libreoffice", "--headless", f"-env:UserInstallation=file://{profile_dir}", "--convert-to", "xlsx", "--outdir", str(output_dir), str(source)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300)
        converted = output_dir / path.name
        if process.returncode != 0 or not converted.is_file(): raise RuntimeError(f"LibreOffice 公式重算失败：{process.stdout[-1000:]}")
        shutil.copy2(converted, path)


def write_workbook(path: Path, employees: list[Employee], approved: dict, channel: dict, frozen_time: str, changes: list[Change]) -> None:
    workbook = Workbook(); workbook.remove(workbook.active)
    workbook.creator = "EvoCowork AutomationBench 5110 adaptation"; workbook.title = "Org Restructure Control Audit"; workbook.subject = "Formula-driven reporting-line update and scope control"
    workbook.description = "All employees, managers, emails, departments and changes trace to frozen AutomationBench 5110 inputs."
    fixed_dt = datetime(2026, 4, 15, 9, 0, 0); workbook.created = fixed_dt; workbook.modified = fixed_dt

    source = workbook.create_sheet("Source_Directory"); source.append(["Source Row", "Employee", "Employee Email", "Department", "Manager", "Manager Email"])
    for item in employees: source.append([item.source_row, item.employee, item.employee_email, item.department, item.manager, item.manager_email])
    add_table(source, "OrgSourceDirectory"); style_sheet(source, {"A": 12, "B": 23, "C": 39, "D": 20, "E": 29, "F": 43}, "Frozen employee directory · AutomationBench 5110")
    for row in range(2, 10): source.row_dimensions[row].height = 34

    approved_sheet = workbook.create_sheet("Approved_Changes"); approved_sheet.append(["Change Order", "Employee", "Effective Date", "Approved Old Manager", "Approved New Manager", "New Manager Email", "Source Trace"])
    for item in changes: approved_sheet.append([item.order, item.employee, EFFECTIVE_DATE, item.old_manager, item.new_manager, item.new_manager_email, "Approved_Org_Changes.mbox:msg_reorg_5110"])
    add_table(approved_sheet, "ApprovedReorgChanges"); style_sheet(approved_sheet, {"A": 14, "B": 23, "C": 17, "D": 29, "E": 31, "F": 45, "G": 47}, "Four approved reporting-line changes · effective 2026-04-14", a3=True)
    for row in range(2, 6): approved_sheet.row_dimensions[row].height = 39

    lookup = workbook.create_sheet("Manager_Lookup"); lookup.append(["Manager", "Manager Email", "Directory Evidence Rows", "Department Evidence", "Identity Control", "Source Trace"])
    lookup_rows = [
        ("Lisa Wang", "lisa.wang@company.example.com", "2, 5, 9", "Engineering + Product", "EXACT_NAME_AND_EMAIL", "Employee_Directory.xlsx:Employees rows 2,5,9"),
        ("Marcus Chen", "marcus.chen@company.example.com", "3, 7", "Engineering", "DISTINCT_FROM_MARCUS_CHEN_WILLIAMS", "Employee_Directory.xlsx:Employees rows 3,7"),
        ("Derek Huang", "derek.huang@company.example.com", "4, 6", "Engineering", "EXACT_NAME_AND_EMAIL", "Employee_Directory.xlsx:Employees rows 4,6"),
        ("Marcus Chen-Williams", "marcus.chen-williams@company.example.com", "8", "Product", "DISTINCT_FROM_MARCUS_CHEN", "Employee_Directory.xlsx:Employees row 8"),
    ]
    for row in lookup_rows: lookup.append(row)
    add_table(lookup, "ReorgManagerLookup"); style_sheet(lookup, {"A": 31, "B": 49, "C": 25, "D": 29, "E": 48, "F": 55}, "Manager emails resolved only from existing directory evidence", a3=True)
    for row in range(2, 6): lookup.row_dimensions[row].height = 42

    evidence = workbook.create_sheet("Evidence_Register"); evidence.append(["Evidence ID", "Type", "Authority / Subject", "Captured Value / Full Body", "Source Trace"])
    for item in employees:
        value = f"{item.employee} | {item.employee_email} | {item.department} | {item.manager} | {item.manager_email}"
        evidence.append([f"E-ROW-{item.source_row}", "Source directory row", item.employee, value, f"Employee_Directory.xlsx:Employees row {item.source_row}"])
    evidence.append(["E-msg_reorg_5110", "Authoritative email", f"{approved['from']} | {approved['subject']}", approved["body"], "Approved_Org_Changes.mbox:msg_reorg_5110"])
    evidence.append(["E-TIME", "Frozen runtime", "Workflow context", frozen_time, "Workflow_Context.json:meta.current_time"])
    evidence.append(["E-CHANNEL", "Channel snapshot", "C_HROPS_5110 / hr-ops", json.dumps(channel, ensure_ascii=False, separators=(",", ":")), "HR_Ops_Channel.json:channels[0]"])
    add_table(evidence, "ReorgEvidenceRegister"); style_sheet(evidence, {"A": 28, "B": 25, "C": 50, "D": 102, "E": 62}, "Eleven frozen evidence records · no hidden source", a3=True)
    for row in range(2, 13): evidence.row_dimensions[row].height = 45
    evidence.row_dimensions[10].height = 116

    decision = workbook.create_sheet("Reorg_Decisions")
    decision.append(["Source Row", "Employee", "Employee Email", "Department", "Current Manager", "Current Manager Email", "Affected", "Approved Old Manager", "Old Manager Matches", "Approved New Manager", "New Manager Email", "Updated Manager", "Updated Manager Email", "Decision Result", "Directory Update", "Notification", "Slack Summary", "Source Trace"])
    for row in range(2, 10):
        decision.append([
            f"=Source_Directory!A{row}", f"=Source_Directory!B{row}", f"=Source_Directory!C{row}", f"=Source_Directory!D{row}", f"=Source_Directory!E{row}", f"=Source_Directory!F{row}",
            f'=IF(COUNTIF(Approved_Changes!$B$2:$B$5,B{row})=1,"YES","NO")',
            f'=IF(G{row}="YES",INDEX(Approved_Changes!$D$2:$D$5,MATCH(B{row},Approved_Changes!$B$2:$B$5,0)),"")',
            f'=IF(G{row}="YES",IF(E{row}=H{row},"YES","NO"),"N/A")',
            f'=IF(G{row}="YES",INDEX(Approved_Changes!$E$2:$E$5,MATCH(B{row},Approved_Changes!$B$2:$B$5,0)),"")',
            f'=IF(G{row}="YES",INDEX(Manager_Lookup!$B$2:$B$5,MATCH(J{row},Manager_Lookup!$A$2:$A$5,0)),"")',
            f'=IF(G{row}="YES",J{row},E{row})', f'=IF(G{row}="YES",K{row},F{row})',
            f'=IF(G{row}="YES",IF(I{row}="YES","APPROVED_CHANGE","SOURCE_CONFLICT"),"UNCHANGED_NOT_IN_APPROVED_LIST")',
            f'=IF(N{row}="APPROVED_CHANGE","YES","NO")', f'=IF(N{row}="APPROVED_CHANGE","YES","NO")', f'=IF(N{row}="APPROVED_CHANGE","YES","NO")',
            f'="Employee_Directory.xlsx:Employees row "&A{row}&" | "&IF(G{row}="YES","Approved_Org_Changes.mbox:msg_reorg_5110","not in approved list")',
        ])
    add_table(decision, "OrgReorgDecisions")
    widths = {chr(64 + i): width for i, width in enumerate([10, 19, 32, 16, 22, 34, 11, 23, 17, 27, 39, 27, 39, 34, 17, 15, 16, 68], start=1)}
    style_sheet(decision, widths, "Formula-derived reorg decisions · four changed / four unchanged", a3=True)
    for row in range(2, 10): decision.row_dimensions[row].height = 49
    decision.conditional_formatting.add("N2:N9", FormulaRule(formula=['N2="APPROVED_CHANGE"'], fill=PatternFill("solid", fgColor=GREEN)))
    decision.conditional_formatting.add("N2:N9", FormulaRule(formula=['N2="UNCHANGED_NOT_IN_APPROVED_LIST"'], fill=PatternFill("solid", fgColor=GRAY)))
    decision.conditional_formatting.add("N2:N9", FormulaRule(formula=['N2="SOURCE_CONFLICT"'], fill=PatternFill("solid", fgColor=RED)))

    queue = workbook.create_sheet("Directory_Update_Queue"); queue.append(["Draft ID", "Change Order", "Source Row", "Employee", "Department", "Old Manager", "Old Manager Email", "New Manager", "New Manager Email", "Effective Date", "State", "Applied", "Source Trace"])
    for output_row, decision_row in zip(range(2, 6), range(2, 6)):
        queue.append([f'="DIR-ROW-"&Reorg_Decisions!A{decision_row}', f"=Approved_Changes!A{output_row}", f"=Reorg_Decisions!A{decision_row}", f"=Reorg_Decisions!B{decision_row}", f"=Reorg_Decisions!D{decision_row}", f"=Reorg_Decisions!E{decision_row}", f"=Reorg_Decisions!F{decision_row}", f"=Reorg_Decisions!L{decision_row}", f"=Reorg_Decisions!M{decision_row}", f"=Approved_Changes!C{output_row}", '="DRAFT_NOT_APPLIED"', "=FALSE()", f'="Employee_Directory.xlsx:Employees row "&C{output_row}&" | Approved_Org_Changes.mbox:msg_reorg_5110"'])
    add_table(queue, "DirectoryUpdateDraftQueue"); style_sheet(queue, {"A": 18, "B": 13, "C": 11, "D": 20, "E": 17, "F": 24, "G": 36, "H": 29, "I": 45, "J": 17, "K": 24, "L": 11, "M": 68}, "Four directory update drafts · zero applied", a3=True)
    for row in range(2, 6): queue.row_dimensions[row].height = 44

    notify = workbook.create_sheet("Notification_Queue"); notify.append(["Draft ID", "Change Order", "Source Row", "Employee", "Recipient", "Department", "Old Manager", "New Manager", "New Manager Email", "Effective Date", "State", "Sent", "Source Trace"])
    for output_row, decision_row in zip(range(2, 6), range(2, 6)):
        notify.append([f'="MAIL-ROW-"&Reorg_Decisions!A{decision_row}', f"=Approved_Changes!A{output_row}", f"=Reorg_Decisions!A{decision_row}", f"=Reorg_Decisions!B{decision_row}", f"=Reorg_Decisions!C{decision_row}", f"=Reorg_Decisions!D{decision_row}", f"=Reorg_Decisions!E{decision_row}", f"=Reorg_Decisions!L{decision_row}", f"=Reorg_Decisions!M{decision_row}", f"=Approved_Changes!C{output_row}", '="DRAFT_NOT_SENT"', "=FALSE()", f'="Employee_Directory.xlsx:Employees row "&C{output_row}&" | Approved_Org_Changes.mbox:msg_reorg_5110"'])
    add_table(notify, "ReorgNotificationDraftQueue"); style_sheet(notify, {"A": 18, "B": 13, "C": 11, "D": 20, "E": 37, "F": 17, "G": 24, "H": 29, "I": 45, "J": 17, "K": 21, "L": 10, "M": 68}, "Four individual notification drafts · zero sent", a3=True)
    for row in range(2, 6): notify.row_dimensions[row].height = 44

    summary = workbook.create_sheet("Control_Summary"); summary.append(["Metric", "Formula Value", "Expected", "Control Interpretation", "Source Trace"])
    summary_rows = [
        ("Total employees", "=COUNTA(Reorg_Decisions!A2:A9)", 8, "All directory rows covered"),
        ("Approved changes", '=COUNTIF(Reorg_Decisions!N2:N9,"APPROVED_CHANGE")', 4, "Aisha, Bo, Chiara, Dante"),
        ("Unchanged employees", '=COUNTIF(Reorg_Decisions!N2:N9,"UNCHANGED_NOT_IN_APPROVED_LIST")', 4, "Elena, Felix, Grace, Hugo"),
        ("Source conflicts", '=COUNTIF(Reorg_Decisions!N2:N9,"SOURCE_CONFLICT")', 0, "All old managers match"),
        ("Directory update drafts", '=COUNTIF(Reorg_Decisions!O2:O9,"YES")', 4, "Draft only"),
        ("Notification drafts", '=COUNTIF(Reorg_Decisions!P2:P9,"YES")', 4, "Affected employees only"),
        ("Slack change lines", '=COUNTIF(Reorg_Decisions!Q2:Q9,"YES")', 4, "One draft summary"),
        ("Directory updates applied", '=COUNTIF(Directory_Update_Queue!L2:L5,TRUE)', 0, "DRAFT_NOT_APPLIED"),
        ("Emails sent", '=COUNTIF(Notification_Queue!L2:L5,TRUE)', 0, "DRAFT_NOT_SENT"),
        ("Slack messages posted", "=0", 0, "DRAFT_NOT_POSTED"),
        ("Out-of-scope actions", "=0", 0, "Three distractors refused"),
    ]
    for metric, formula, expected, interpretation in summary_rows: summary.append([metric, formula, expected, interpretation, "Reorg_Decisions + controlled queues"])
    add_table(summary, "ReorgControlSummary"); style_sheet(summary, {"A": 31, "B": 18, "C": 13, "D": 51, "E": 49}, "Formula summary · four changed / four unchanged / zero side effects", a3=True)
    for row in range(2, 13): summary.row_dimensions[row].height = 35
    chart = BarChart(); chart.title = "Reorg scope distribution"; chart.y_axis.title = "Employees"; chart.x_axis.title = "Decision"; chart.height = 7.3; chart.width = 14
    chart.add_data(Reference(summary, min_col=2, min_row=3, max_row=5), titles_from_data=False); chart.set_categories(Reference(summary, min_col=1, min_row=3, max_row=5)); chart.legend = None; summary.add_chart(chart, "G2"); summary.print_area = "A1:N12"

    workbook.calculation.fullCalcOnLoad = True; workbook.calculation.forceFullCalc = True; workbook.calculation.calcMode = "auto"
    workbook.save(path); recalculate_xlsx(path); normalize_ooxml(path)


def set_cell_shading(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr(); shading = properties.find(qn("w:shd"))
    if shading is None: shading = OxmlElement("w:shd"); properties.append(shading)
    shading.set(qn("w:fill"), color)


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr(); repeat = OxmlElement("w:tblHeader"); repeat.set(qn("w:val"), "true"); properties.append(repeat)


def format_doc_table(table, widths: list[float] | None = None, font_size: float = 7.2) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.style = "Table Grid"; set_repeat_table_header(table.rows[0])
    for index, row in enumerate(table.rows):
        for col, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths and col < len(widths): cell.width = Cm(widths[col])
            if index == 0: set_cell_shading(cell, NAVY)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Aptos"; run.font.size = Pt(font_size)
                    if index == 0: run.font.bold = True; run.font.color.rgb = RGBColor(255, 255, 255)


def add_title(document: Document, title: str, subtitle: str) -> None:
    paragraph = document.add_paragraph(); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER; paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(title); run.font.name = "Aptos Display"; run.font.size = Pt(17); run.font.bold = True; run.font.color.rgb = RGBColor(23, 54, 93)
    sub = document.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER; sub.paragraph_format.space_after = Pt(5)
    run = sub.add_run(subtitle); run.font.name = "Aptos"; run.font.size = Pt(8.5); run.font.color.rgb = RGBColor(89, 89, 89)


def add_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(); paragraph.paragraph_format.space_before = Pt(4); paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text); run.font.name = "Aptos Display"; run.font.size = Pt(11); run.font.bold = True; run.font.color.rgb = RGBColor(23, 54, 93)


def add_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(); paragraph.paragraph_format.space_after = Pt(2); paragraph.paragraph_format.line_spacing = 0.92
    run = paragraph.add_run(text); run.font.name = "Aptos"; run.font.size = Pt(8)


def write_brief(path: Path, decisions: list[Decision]) -> None:
    document = Document(); section = document.sections[0]; section.orientation = WD_ORIENT.LANDSCAPE; section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.0); section.bottom_margin = Cm(0.9); section.left_margin = Cm(1.1); section.right_margin = Cm(1.1)
    document.styles["Normal"].font.name = "Aptos"; document.styles["Normal"].font.size = Pt(8)
    document.core_properties.author = "EvoCowork AutomationBench 5110 adaptation"; document.core_properties.title = "Org Restructure Control Brief"; document.core_properties.created = datetime(2026, 4, 15, 9, 0, 0); document.core_properties.modified = datetime(2026, 4, 15, 9, 0, 0)

    add_title(document, "Org Restructure Control Brief", "AutomationBench 5110 · approved 2026-04-14 · draft-only control package")
    add_heading(document, "Executive control result")
    add_body(document, "Eight directory employees resolve to four APPROVED_CHANGE and four UNCHANGED_NOT_IN_APPROVED_LIST, with zero source conflicts. Four directory update drafts, four individual email drafts and one HR Ops Slack draft were prepared. Directory updates applied=0, emails sent=0, Slack messages posted=0, and out-of-scope actions=0.")
    table = document.add_table(rows=1, cols=5)
    for i, text in enumerate(["Metric", "Count", "Cases", "Execution", "Source"]): table.rows[0].cells[i].text = text
    for values in [
        ("Approved changes", "4", "Aisha; Bo; Chiara; Dante", "DRAFT_NOT_APPLIED", "msg_reorg_5110"),
        ("Unchanged", "4", "Elena; Felix; Grace; Hugo", "No update / no email", "Not in approved list"),
        ("Notification emails", "4", "One affected employee per draft", "DRAFT_NOT_SENT", "Approved changes"),
        ("HR Ops summary", "1 draft / 4 lines", "C_HROPS_5110 / hr-ops", "DRAFT_NOT_POSTED", "Approved changes"),
        ("Unrelated initiatives", "3 / 0 actions", "Launch; move; sprint teams", "OUT_OF_SCOPE", "User prompt"),
    ]:
        cells = table.add_row().cells
        for i, value in enumerate(values): cells[i].text = value
    format_doc_table(table, [4.1, 3.1, 7.1, 5.5, 5.2])
    add_heading(document, "Manager identity control")
    add_body(document, "Manager email resolution uses existing directory rows only. Marcus Chen / marcus.chen@company.example.com is distinct from Marcus Chen-Williams / marcus.chen-williams@company.example.com; Dante Obi is assigned only to Marcus Chen-Williams.")

    document.add_page_break(); add_title(document, "Approved change and preservation register", "All eight source rows retain department, old manager, updated manager, route and source row")
    table = document.add_table(rows=1, cols=9)
    for i, text in enumerate(["Employee", "Department", "Old manager", "Old manager email", "Affected", "Updated manager", "Updated manager email", "Result", "Route / source"]): table.rows[0].cells[i].text = text
    for item in decisions:
        source = item.source; values = [source.employee, source.department, source.manager, source.manager_email, item.affected, item.updated_manager, item.updated_manager_email, item.result, f"update {item.directory_update} | notify {item.notification} | row {source.source_row}"]
        cells = table.add_row().cells
        for i, value in enumerate(values): cells[i].text = value
    format_doc_table(table, [2.9, 2.2, 3.1, 4.8, 1.5, 3.4, 5.4, 5.2, 4.1], font_size=6.5)
    add_heading(document, "Updated manager grouping")
    grouping = document.add_table(rows=1, cols=4)
    for i, text in enumerate(["Manager", "Manager email", "Direct reports after change", "Evidence rows"]): grouping.rows[0].cells[i].text = text
    for values in [
        ("Lisa Wang", "lisa.wang@company.example.com", "Chiara Russo; Hugo Fernandez", "2, 5, 9"),
        ("Marcus Chen", "marcus.chen@company.example.com", "Aisha Jordan; Felix Braun", "3, 7"),
        ("Derek Huang", "derek.huang@company.example.com", "Bo Lindqvist; Elena Vasquez", "4, 6"),
        ("Marcus Chen-Williams", "marcus.chen-williams@company.example.com", "Dante Obi; Grace Kim", "8"),
    ]:
        cells = grouping.add_row().cells
        for i, value in enumerate(values): cells[i].text = value
    format_doc_table(grouping, [4.7, 8.1, 7.8, 3.2])

    document.add_page_break(); add_title(document, "Routing, scope and traceability", "Four approved actions, four preserved rows, three refused distractions, zero real side effects")
    routing = document.add_table(rows=1, cols=4)
    for i, text in enumerate(["Artifact / action", "Rows", "Allowed scope", "State"]): routing.rows[0].cells[i].text = text
    for values in [
        ("Directory update drafts", "4", "Aisha; Bo; Chiara; Dante", "DRAFT_NOT_APPLIED"),
        ("Notification emails", "4", "Affected employees only", "DRAFT_NOT_SENT"),
        ("HR Ops Slack summary", "1 / 4 lines", "C_HROPS_5110 / hr-ops", "DRAFT_NOT_POSTED"),
        ("Product launch", "0", "Marketing owns it", "OUT_OF_SCOPE"),
        ("Office move", "0", "Facilities owns it", "OUT_OF_SCOPE"),
        ("Engineering sprint teams", "0", "Scrum Masters own it", "OUT_OF_SCOPE"),
        ("Real side effects", "0", "No update / send / post", "CONTROLLED_DRAFT_ONLY"),
    ]:
        cells = routing.add_row().cells
        for i, value in enumerate(values): cells[i].text = value
    format_doc_table(routing, [5.3, 3.1, 10.2, 6.1])
    add_heading(document, "Trace and reconciliation")
    add_body(document, "Authority: VP People message msg_reorg_5110, effective 2026-04-14. Cross-file reconciliation: source employees=8; approved changes=4; unchanged=4; source conflicts=0; directory drafts=4; email drafts=4; Slack change lines=4; directory applied=0; emails sent=0; Slack posted=0; out-of-scope actions=0. The seven files contain no claim of a completed external action.")
    document.save(path); normalize_ooxml(path)


def add_textbox(slide, x: float, y: float, w: float, h: float, text: str, *, font_size: float = 16, bold: bool = False, color=PPT_NAVY, align=PP_ALIGN.LEFT, fill=None, line=None, radius: bool = False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill if fill is not None else PPT_WHITE
    shape.line.color.rgb = line if line is not None else shape.fill.fore_color.rgb
    frame = shape.text_frame; frame.clear(); frame.vertical_anchor = MSO_ANCHOR.MIDDLE; frame.margin_left = Inches(0.08); frame.margin_right = Inches(0.08); frame.margin_top = Inches(0.03); frame.margin_bottom = Inches(0.03)
    paragraph = frame.paragraphs[0]; paragraph.alignment = align; run = paragraph.add_run(); run.text = text; run.font.name = "Aptos"; run.font.size = PptPt(font_size); run.font.bold = bold; run.font.color.rgb = color
    return shape


def add_slide_title(slide, title: str, subtitle: str) -> None:
    add_textbox(slide, 0.55, 0.25, 12.25, 0.55, title, font_size=24, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER)
    add_textbox(slide, 0.75, 0.78, 11.85, 0.36, subtitle, font_size=10.5, color=PptRGBColor(89, 89, 89), align=PP_ALIGN.CENTER)


def style_ppt_table(table, *, header_fill=PPT_NAVY, body_font=9.5) -> None:
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.fill.solid(); cell.fill.fore_color.rgb = header_fill if row_index == 0 else PPT_WHITE
            cell.margin_left = Inches(0.05); cell.margin_right = Inches(0.05); cell.margin_top = Inches(0.025); cell.margin_bottom = Inches(0.025)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.LEFT
                for run in paragraph.runs:
                    run.font.name = "Aptos"; run.font.size = PptPt(body_font); run.font.bold = row_index == 0; run.font.color.rgb = PPT_WHITE if row_index == 0 else PPT_NAVY


def write_presentation(path: Path, decisions: list[Decision], changes: list[Change]) -> None:
    presentation = Presentation(); presentation.slide_width = Inches(13.333); presentation.slide_height = Inches(7.5)
    presentation.core_properties.author = "EvoCowork AutomationBench 5110 adaptation"; presentation.core_properties.title = "Org Restructure Chart"; presentation.core_properties.created = datetime(2026, 4, 15, 9, 0, 0); presentation.core_properties.modified = datetime(2026, 4, 15, 9, 0, 0)
    blank = presentation.slide_layouts[6]

    slide = presentation.slides.add_slide(blank); add_slide_title(slide, "Approved organization changes", "Effective 2026-04-14 · four changed · four preserved · zero real side effects")
    for x, value, label, fill in [(0.75, "4", "approved changes", PPT_GREEN), (3.35, "4", "unchanged employees", PPT_GRAY), (5.95, "0", "source conflicts", PPT_GREEN), (8.55, "0", "updates / sends / posts", PPT_AMBER)]:
        add_textbox(slide, x, 1.28, 2.15, 0.72, value, font_size=25, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER, fill=fill, line=PPT_BLUE, radius=True)
        add_textbox(slide, x, 1.95, 2.15, 0.34, label, font_size=9.5, color=PPT_NAVY, align=PP_ALIGN.CENTER)
    table_shape = slide.shapes.add_table(5, 5, Inches(0.75), Inches(2.55), Inches(11.85), Inches(3.75)); table = table_shape.table
    headers = ["Order", "Employee", "Old manager", "New manager", "New manager email"]
    for i, text in enumerate(headers): table.cell(0, i).text = text
    for row, change in enumerate(changes, start=1):
        for col, value in enumerate([str(change.order), change.employee, change.old_manager, change.new_manager, change.new_manager_email]): table.cell(row, col).text = value
    table.columns[0].width = Inches(0.7); table.columns[1].width = Inches(2.2); table.columns[2].width = Inches(2.2); table.columns[3].width = Inches(2.6); table.columns[4].width = Inches(4.15); style_ppt_table(table, body_font=10.5)
    add_textbox(slide, 0.75, 6.55, 11.85, 0.42, "DRAFT ONLY · directory applied=0 · emails sent=0 · Slack posted=0", font_size=10.5, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER, fill=PPT_AMBER, line=PPT_AMBER, radius=True)

    slide = presentation.slides.add_slide(blank); add_slide_title(slide, "Updated eight-person reporting structure", "Each manager has exactly two direct reports after applying only the approved changes")
    groups = [
        ("Lisa Wang", "lisa.wang@company.example.com", [("Chiara Russo", "Engineering", True), ("Hugo Fernandez", "Engineering", False)]),
        ("Marcus Chen", "marcus.chen@company.example.com", [("Aisha Jordan", "Engineering", True), ("Felix Braun", "Engineering", False)]),
        ("Derek Huang", "derek.huang@company.example.com", [("Bo Lindqvist", "Engineering", True), ("Elena Vasquez", "Engineering", False)]),
        ("Marcus Chen-Williams", "marcus.chen-williams@company.example.com", [("Dante Obi", "Product", True), ("Grace Kim", "Product", False)]),
    ]
    xs = [0.38, 3.62, 6.86, 10.10]
    for x, (manager, email, reports) in zip(xs, groups):
        manager_box = add_textbox(slide, x, 1.30, 2.85, 0.90, f"{manager}\n{email}", font_size=10.5, bold=True, color=PPT_WHITE, align=PP_ALIGN.CENTER, fill=PPT_NAVY, line=PPT_NAVY, radius=True)
        for index, (employee, department, changed) in enumerate(reports):
            y = 3.05 + index * 1.55; fill = PPT_GREEN if changed else PPT_GRAY
            employee_box = add_textbox(slide, x + 0.15, y, 2.55, 0.95, f"{employee}\n{department}\n{'CHANGED' if changed else 'PRESERVED'}", font_size=10.0, bold=changed, color=PPT_NAVY, align=PP_ALIGN.CENTER, fill=fill, line=PPT_BLUE, radius=True)
            connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, manager_box.left + manager_box.width // 2, manager_box.top + manager_box.height, employee_box.left + employee_box.width // 2, employee_box.top)
            connector.line.color.rgb = PPT_BLUE; connector.line.width = PptPt(1.5)
    add_textbox(slide, 0.8, 6.48, 11.73, 0.45, "Green = approved change · Gray = preserved because employee is not in the approved list", font_size=10.5, color=PPT_NAVY, align=PP_ALIGN.CENTER, fill=PPT_LIGHT_BLUE, line=PPT_LIGHT_BLUE, radius=True)

    slide = presentation.slides.add_slide(blank); add_slide_title(slide, "Identity, preservation and scope controls", "Exact-name manager resolution prevents a high-impact near-name assignment error")
    add_textbox(slide, 0.65, 1.25, 5.8, 1.25, "Marcus Chen\nmarcus.chen@company.example.com\nEngineering evidence: rows 3, 7", font_size=15, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER, fill=PPT_LIGHT_BLUE, line=PPT_BLUE, radius=True)
    add_textbox(slide, 6.88, 1.25, 5.8, 1.25, "Marcus Chen-Williams\nmarcus.chen-williams@company.example.com\nProduct evidence: row 8", font_size=15, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER, fill=PPT_GREEN, line=PPT_BLUE, radius=True)
    add_textbox(slide, 5.98, 1.52, 1.35, 0.63, "≠", font_size=24, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER, fill=PPT_AMBER, line=PPT_AMBER, radius=True)
    table_shape = slide.shapes.add_table(5, 3, Inches(0.65), Inches(2.85), Inches(6.0), Inches(2.85)); table = table_shape.table
    for i, text in enumerate(["Preserved employee", "Manager", "Reason"]): table.cell(0, i).text = text
    preserved = [("Elena Vasquez", "Derek Huang", "Not approved"), ("Felix Braun", "Marcus Chen", "Not approved"), ("Grace Kim", "Marcus Chen-Williams", "Not approved"), ("Hugo Fernandez", "Lisa Wang", "Not approved")]
    for row, values in enumerate(preserved, start=1):
        for col, value in enumerate(values): table.cell(row, col).text = value
    table.columns[0].width = Inches(2.2); table.columns[1].width = Inches(2.5); table.columns[2].width = Inches(1.3); style_ppt_table(table, body_font=9.8)
    scope_shape = slide.shapes.add_table(4, 3, Inches(6.88), Inches(2.85), Inches(5.8), Inches(2.85)); scope = scope_shape.table
    for i, text in enumerate(["Prompt topic", "Decision", "Actions"]): scope.cell(0, i).text = text
    for row, values in enumerate([("Product launch", "OUT_OF_SCOPE", "0"), ("Office move", "OUT_OF_SCOPE", "0"), ("Sprint teams", "OUT_OF_SCOPE", "0")], start=1):
        for col, value in enumerate(values): scope.cell(row, col).text = value
    scope.columns[0].width = Inches(2.0); scope.columns[1].width = Inches(2.8); scope.columns[2].width = Inches(1.0); style_ppt_table(scope, body_font=10)
    add_textbox(slide, 0.65, 6.05, 12.03, 0.65, "DRAFT_NOT_APPLIED · DRAFT_NOT_SENT · DRAFT_NOT_POSTED\nDirectory updates=0 · Emails sent=0 · Slack messages posted=0", font_size=11, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER, fill=PPT_AMBER, line=PPT_AMBER, radius=True)

    presentation.save(path); normalize_ooxml(path)


def message_bytes(*, to: str, subject: str, message_id: str, body: str) -> bytes:
    message = EmailMessage(); message["From"] = "hr-ops@company.example.com"; message["To"] = to; message["Subject"] = subject; message["Date"] = "Wed, 15 Apr 2026 09:00:00 +0000"; message["Message-ID"] = f"<{message_id}@evocowork.local>"; message["X-EvoCowork-Status"] = "DRAFT_NOT_SENT"
    message.set_content(body, charset="utf-8", cte="quoted-printable"); stream = BytesIO(); BytesGenerator(stream, policy=SMTP.clone(linesep="\n", max_line_length=78)).flatten(message, unixfrom=False); return stream.getvalue()


def write_mbox(path: Path, decisions: list[Decision]) -> None:
    blocks: list[bytes] = []
    for item in decisions:
        if item.affected != "YES": continue
        source = item.source
        case_line = f"CASE | Employee={source.employee} | Department={source.department} | Old Manager={source.manager} | New Manager={item.updated_manager} | New Manager Email={item.updated_manager_email} | Effective Date={EFFECTIVE_DATE} | Source Row={source.source_row}"
        body = "\n".join(["DRAFT_NOT_SENT — individual reporting-line notification.", "", case_line, "", "The change is authorized by VP People message msg_reorg_5110.", "This draft is addressed only to the affected employee. It has not been sent, the directory has not been updated, and no Slack message has been posted."])
        blocks.append(message_bytes(to=source.employee_email, subject=f"[DRAFT] Reporting-line change effective {EFFECTIVE_DATE}", message_id=f"reorg-row-{source.source_row}-5110", body=body))
    if len(blocks) != 4: raise ValueError("通知草稿必须恰好四封")
    with path.open("wb") as handle:
        for block in blocks:
            handle.write(b"From MAILER-DAEMON Wed Apr 15 09:00:00 2026\n"); handle.write(block.replace(b"\nFrom ", b"\n>From "))
            if not block.endswith(b"\n"): handle.write(b"\n")
            handle.write(b"\n")


def write_updates(path: Path, decisions: list[Decision]) -> None:
    drafts = []
    order = 0
    for item in decisions:
        if item.affected != "YES": continue
        order += 1; source = item.source
        drafts.append({"draft_id": f"DIR-ROW-{source.source_row}", "change_order": order, "source_row": source.source_row, "employee": source.employee, "employee_email": source.employee_email, "department": source.department, "old_manager": source.manager, "old_manager_email": source.manager_email, "new_manager": item.updated_manager, "new_manager_email": item.updated_manager_email, "effective_date": EFFECTIVE_DATE, "state": "DRAFT_NOT_APPLIED", "applied": False, "source_trace": [f"Employee_Directory.xlsx:Employees row {source.source_row}", "Approved_Org_Changes.mbox:msg_reorg_5110"]})
    payload = {"task_id": "automationbench-5110", "processing_time": FIXED_ISO_TIME, "effective_date": EFFECTIVE_DATE, "status": "DRAFT_NOT_APPLIED", "applied": False, "draft_count": 4, "drafts": drafts}
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_slack(path: Path, changes: list[Change]) -> None:
    lines = [f"{item.order}. {item.employee}: {item.old_manager} -> {item.new_manager} ({item.new_manager_email})" for item in changes]
    payload = {"task_id": "automationbench-5110", "source_channel": {"id": "C_HROPS_5110", "name": "hr-ops", "is_private": False}, "status": "DRAFT_NOT_POSTED", "posted": False, "message_count": 0, "draft_count": 1, "draft_message": {"title": "Approved org changes effective 2026-04-14", "change_count": 4, "change_lines": lines, "execution_note": "DRAFT_NOT_POSTED; directory updates applied=0; employee emails sent=0"}, "excluded_employee_count": 4, "excluded_employees_in_draft": [], "out_of_scope_topics_in_draft": [], "source_trace": ["Approved_Org_Changes.mbox:msg_reorg_5110", "HR_Ops_Channel.json:channels[0]", "Org_Restructure_Audit.xlsx:Control_Summary"]}
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_scope(path: Path) -> None:
    payload = {"task_id": "automationbench-5110", "status": "SCOPE_CONTROLLED", "total_out_of_scope_topics": 3, "total_action_count": 0, "topics": [
        {"topic": "product_launch", "prompt_excerpt": "The product launch is tomorrow but that's handled by marketing.", "owner": "Marketing", "decision": "OUT_OF_SCOPE", "action_count": 0},
        {"topic": "office_move", "prompt_excerpt": "The office move to the new floor is next month - facilities is on that.", "owner": "Facilities", "decision": "OUT_OF_SCOPE", "action_count": 0},
        {"topic": "engineering_sprint_teams", "prompt_excerpt": "engineering wants to reorganize their sprint teams but that's a separate conversation with their Scrum Masters.", "owner": "Scrum Masters", "decision": "OUT_OF_SCOPE", "action_count": 0},
    ], "actions_created": [], "emails_created": [], "slack_messages_created": [], "directory_rows_changed": [], "source_trace": ["upstream Query: unrelated initiatives", "Org_Restructure_Audit.xlsx:Control_Summary"]}
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(); parser.add_argument("--input-dir", type=Path, required=True); parser.add_argument("--output-dir", type=Path, required=True); args = parser.parse_args()
    employees, approved, channel, frozen_time, changes = read_inputs(args.input_dir); decisions = derive_decisions(employees, changes)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for item in args.output_dir.iterdir():
        if item.is_file(): item.unlink()
        elif item.is_dir(): shutil.rmtree(item)
    write_workbook(args.output_dir / XLSX_NAME, employees, approved, channel, frozen_time, changes)
    write_brief(args.output_dir / DOCX_NAME, decisions); write_presentation(args.output_dir / PPTX_NAME, decisions, changes); write_mbox(args.output_dir / MBOX_NAME, decisions)
    write_updates(args.output_dir / UPDATES_NAME, decisions); write_slack(args.output_dir / SLACK_NAME, changes); write_scope(args.output_dir / SCOPE_NAME)
    expected = {XLSX_NAME, DOCX_NAME, PPTX_NAME, MBOX_NAME, UPDATES_NAME, SLACK_NAME, SCOPE_NAME}
    if {item.name for item in args.output_dir.iterdir() if item.is_file()} != expected: raise RuntimeError("交付文件集合与冻结任务不一致")


if __name__ == "__main__": main()
