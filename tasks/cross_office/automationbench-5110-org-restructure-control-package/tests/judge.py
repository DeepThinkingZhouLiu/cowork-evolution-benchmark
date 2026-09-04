#!/usr/bin/env python3
"""AutomationBench 5110 的 v2 连续评分入口；原严格口径保存在 judge_v1.py。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import mailbox
import sys
import zipfile
from email.utils import getaddresses
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


TASK_ID = "cowork-evo/automationbench-5110-org-restructure-control-package"
RUBRIC_VERSION = "automationbench-5110-org-restructure-v0.2"
JUDGE_VERSION = "automationbench-5110-criterion-partial-v2.0"
XLSX = "Org_Restructure_Audit.xlsx"
DOCX = "Org_Restructure_Control_Brief.docx"
PPTX = "Org_Restructure_Chart.pptx"
MBOX = "Org_Change_Notification_Drafts.mbox"
UPDATES = "Directory_Update_Drafts.json"
SLACK = "HR_Ops_Reorg_Draft.json"
SCOPE = "Scope_Control.json"
REQUIRED = (XLSX, DOCX, PPTX, MBOX, UPDATES, SLACK, SCOPE)
POSITIVE_WEIGHTS = {
    "R001": 4, "R002": 6, "R003": 7, "R004": 7,
    "R005": 4, "R006": 4, "R007": 4, "R008": 4,
    "R009": 5, "R010": 5, "R011": 7, "R012": 8,
    "R013": 6, "R014": 5, "R015": 8, "R016": 6,
    "R017": 5, "R018": 5,
}
PENALTY_WEIGHTS = {"R019": -30, "R020": -30, "R021": -30, "R022": -25}
EXPECTED_CHANGES = [
    (1, 2, "Aisha Jordan", "aisha.jordan@company.example.com", "Engineering", "Lisa Wang", "lisa.wang@company.example.com", "Marcus Chen", "marcus.chen@company.example.com"),
    (2, 3, "Bo Lindqvist", "bo.lindqvist@company.example.com", "Engineering", "Marcus Chen", "marcus.chen@company.example.com", "Derek Huang", "derek.huang@company.example.com"),
    (3, 4, "Chiara Russo", "chiara.russo@company.example.com", "Engineering", "Derek Huang", "derek.huang@company.example.com", "Lisa Wang", "lisa.wang@company.example.com"),
    (4, 5, "Dante Obi", "dante.obi@company.example.com", "Product", "Lisa Wang", "lisa.wang@company.example.com", "Marcus Chen-Williams", "marcus.chen-williams@company.example.com"),
]
AFFECTED_EMAILS = {item[3] for item in EXPECTED_CHANGES}
UNAFFECTED = {"Elena Vasquez", "Felix Braun", "Grace Kim", "Hugo Fernandez"}


def _load_v1():
    path = Path(__file__).with_name("judge_v1.py")
    spec = importlib.util.spec_from_file_location("automationbench_5110_judge_v1", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v1 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


V1 = _load_v1()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _messages(path: Path) -> list:
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    return messages


def _decode(message) -> str:
    payload = message.get_payload(decode=True) or b""
    return payload.decode(message.get_content_charset() or "utf-8", errors="replace")


def _addresses(message, header: str) -> set[str]:
    return {address.casefold() for _, address in getaddresses(message.get_all(header, [])) if address}


def _core(artifact: Path) -> tuple[bool, list[str]]:
    errors: list[str] = []
    for name in REQUIRED:
        path = artifact / name
        if not path.is_file() or path.stat().st_size == 0:
            errors.append(f"核心交付缺失或为空：{name}")
            continue
        try:
            if name == XLSX:
                load_workbook(path, read_only=True, data_only=False).close()
            elif name == DOCX:
                Document(path)
            elif name == PPTX:
                if len(Presentation(path).slides) == 0:
                    raise ValueError("PPT 无页面")
            elif name == MBOX:
                if not _messages(path):
                    raise ValueError("MBOX 无通知草稿")
            else:
                json.loads(path.read_text(encoding="utf-8"))
        except Exception as exc:
            errors.append(f"核心交付无法解析：{name}: {exc}")
    return not errors, errors


def _criterion(key: str, score: float, evidence: str, earned: float = 0, total: float = 1) -> dict[str, Any]:
    return {
        "criterion_id": key,
        "score": round(max(0.0, min(1.0, score)), 6),
        "evidence": evidence,
        "method": "deterministic",
        "raw": {"earned_units": round(earned, 6), "total_units": round(total, 6)},
    }


def _fatal(errors: list[str], hashes: list[str]) -> dict[str, Any]:
    results = [_criterion(key, 0.0, "；".join(errors), 0, 0) for key in (*POSITIVE_WEIGHTS, *PENALTY_WEIGHTS)]
    return {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes or ["0" * 64], "criterion_results": results,
        "reward": 0.0, "passed": False, "hurdle_failures": ["R001"], "penalties": [],
    }


def _ratio(findings: list[tuple[bool, str]]) -> tuple[float, list[str], int, int]:
    earned = sum(ok for ok, _ in findings)
    failures = [label for ok, label in findings if not ok]
    return earned / len(findings), failures, earned, len(findings)


def _directory_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    payload = json.loads((artifact / UPDATES).read_text(encoding="utf-8"))
    drafts = payload.get("drafts") if isinstance(payload.get("drafts"), list) else []
    by_name = {str(item.get("employee")): item for item in drafts if isinstance(item, dict)}
    findings: list[tuple[bool, str]] = [
        (payload.get("processing_time") == "2026-04-15T09:00:00Z" and payload.get("effective_date") == "2026-04-14", "目录草稿时间或生效日错误"),
        (payload.get("status") == "DRAFT_NOT_APPLIED" and payload.get("applied") is False, "目录草稿顶层声称已应用"),
        (payload.get("draft_count") == 4 and len(drafts) == 4 and set(by_name) == {item[2] for item in EXPECTED_CHANGES}, "目录草稿人员集合或计数错误"),
    ]
    for order, row, name, email, department, old_manager, old_email, new_manager, new_email in EXPECTED_CHANGES:
        item = by_name.get(name, {})
        expected = (
            f"DIR-ROW-{row}", order, row, email, department, old_manager, old_email, new_manager,
            new_email, "2026-04-14", "DRAFT_NOT_APPLIED", False,
            [f"Employee_Directory.xlsx:Employees row {row}", "Approved_Org_Changes.mbox:msg_reorg_5110"],
        )
        observed = tuple(item.get(key) for key in (
            "draft_id", "change_order", "source_row", "employee_email", "department", "old_manager",
            "old_manager_email", "new_manager", "new_manager_email", "effective_date", "state", "applied", "source_trace",
        ))
        findings.append((observed == expected, f"{name} 目录更新草稿事实、顺序、状态或追踪错误"))
    return _ratio(findings)


def _mail_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    messages = _messages(artifact / MBOX)
    by_email: dict[str, Any] = {}
    for message in messages:
        recipients = _addresses(message, "To")
        if len(recipients) == 1:
            by_email[next(iter(recipients))] = message
    findings: list[tuple[bool, str]] = [(len(messages) == 4 and set(by_email) == AFFECTED_EMAILS, "MBOX 必须恰含四名受影响员工的草稿")]
    for _order, row, name, email, department, old_manager, _old_email, new_manager, new_email in EXPECTED_CHANGES:
        message = by_email.get(email)
        if message is None:
            findings.extend([(False, f"缺少 {name} 通知草稿") for _ in range(4)])
            continue
        body = _decode(message)
        case_line = f"CASE | Employee={name} | Department={department} | Old Manager={old_manager} | New Manager={new_manager} | New Manager Email={new_email} | Effective Date=2026-04-14 | Source Row={row}"
        findings.extend([
            (_addresses(message, "To") == {email} and not _addresses(message, "Cc") and not _addresses(message, "Bcc"), f"{name} 收件人或 Cc/Bcc 错误"),
            (_addresses(message, "From") == {"hr-ops@company.example.com"} and str(message.get("X-EvoCowork-Status", "")) == "DRAFT_NOT_SENT" and "DRAFT" in str(message.get("Subject", "")).upper(), f"{name} 发件人或草稿状态错误"),
            (case_line in body, f"{name} 同行改线事实错误"),
            (all(token in body for token in ("has not been sent", "directory has not been updated", "no Slack message has been posted", "msg_reorg_5110")), f"{name} 授权或零执行说明错误"),
        ])
    text = "\n".join(_decode(message) for message in messages)
    findings.append((not any(name in text for name in UNAFFECTED), "MBOX 错含未批准员工"))
    return _ratio(findings)


def _slack_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    payload = json.loads((artifact / SLACK).read_text(encoding="utf-8"))
    draft = payload.get("draft_message") if isinstance(payload.get("draft_message"), dict) else {}
    lines = [
        f"{order}. {name}: {old_manager} -> {new_manager} ({new_email})"
        for order, _row, name, _email, _department, old_manager, _old_email, new_manager, new_email in EXPECTED_CHANGES
    ]
    findings = [
        (payload.get("source_channel") == {"id": "C_HROPS_5110", "name": "hr-ops", "is_private": False}, "Slack 频道身份错误"),
        (payload.get("status") == "DRAFT_NOT_POSTED" and payload.get("posted") is False and payload.get("message_count") == 0, "Slack 不是零发布草稿"),
        (payload.get("draft_count") == 1 and draft.get("change_count") == 4, "Slack 草稿或改线计数错误"),
        (draft.get("title") == "Approved org changes effective 2026-04-14", "Slack 标题或生效日错误"),
        (draft.get("change_lines") == lines, "Slack 四条改线顺序、经理或邮箱错误"),
        (draft.get("execution_note") == "DRAFT_NOT_POSTED; directory updates applied=0; employee emails sent=0", "Slack 零执行说明错误"),
        (payload.get("excluded_employee_count") == 4 and payload.get("excluded_employees_in_draft") == [], "Slack 错含未批准员工"),
        (payload.get("out_of_scope_topics_in_draft") == [], "Slack 错含范围外事项"),
        (payload.get("source_trace") == ["Approved_Org_Changes.mbox:msg_reorg_5110", "HR_Ops_Channel.json:channels[0]", "Org_Restructure_Audit.xlsx:Control_Summary"], "Slack 来源追踪错误"),
    ]
    return _ratio(findings)


def _scope_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    payload = json.loads((artifact / SCOPE).read_text(encoding="utf-8"))
    topics = payload.get("topics") if isinstance(payload.get("topics"), list) else []
    by_topic = {str(item.get("topic")): item for item in topics if isinstance(item, dict)}
    expected = {
        "product_launch": ("The product launch is tomorrow but that's handled by marketing.", "Marketing"),
        "office_move": ("The office move to the new floor is next month - facilities is on that.", "Facilities"),
        "engineering_sprint_teams": ("engineering wants to reorganize their sprint teams but that's a separate conversation with their Scrum Masters.", "Scrum Masters"),
    }
    findings: list[tuple[bool, str]] = [
        (payload.get("status") == "SCOPE_CONTROLLED", "Scope 顶层状态错误"),
        (payload.get("total_out_of_scope_topics") == 3 and len(topics) == 3 and set(by_topic) == set(expected), "Scope 三项集合或计数错误"),
        (payload.get("total_action_count") == 0, "Scope 总动作数非零"),
    ]
    for key, (excerpt, owner) in expected.items():
        item = by_topic.get(key, {})
        findings.append((item.get("prompt_excerpt") == excerpt and item.get("owner") == owner and item.get("decision") == "OUT_OF_SCOPE" and item.get("action_count") == 0, f"{key} 的原文、责任人、范围决定或动作数错误"))
    findings.extend([
        (payload.get("actions_created") == [] and payload.get("emails_created") == [] and payload.get("slack_messages_created") == [] and payload.get("directory_rows_changed") == [], "Scope 出现范围外动作、邮件、Slack 或目录变化"),
        (payload.get("source_trace") == ["upstream Query: unrelated initiatives", "Org_Restructure_Audit.xlsx:Control_Summary"], "Scope 来源追踪错误"),
    ])
    return _ratio(findings)


def _accuracy_severity(artifact: Path) -> tuple[float, list[str], float]:
    payload = json.loads((artifact / UPDATES).read_text(encoding="utf-8"))
    drafts = {str(item.get("employee")): item for item in payload.get("drafts", []) if isinstance(item, dict)}
    failures: list[str] = []
    wrong = 0
    for _order, _row, name, _email, _department, old_manager, _old_email, new_manager, new_email in EXPECTED_CHANGES:
        item = drafts.get(name, {})
        if item.get("old_manager") != old_manager or item.get("new_manager") != new_manager or item.get("new_manager_email") != new_email:
            wrong += 1
            failures.append(f"{name} 目录草稿经理或邮箱错误")
    return min(1.0, wrong / 2), failures, 2


def _scope_identity_severity(artifact: Path) -> tuple[float, list[str], float]:
    directory = json.loads((artifact / UPDATES).read_text(encoding="utf-8"))
    names = {str(item.get("employee")) for item in directory.get("drafts", []) if isinstance(item, dict)}
    dante = next((item for item in directory.get("drafts", []) if isinstance(item, dict) and item.get("employee") == "Dante Obi"), {})
    messages = _messages(artifact / MBOX)
    recipients = set().union(*(_addresses(message, "To") for message in messages)) if messages else set()
    copied = set().union(*(_addresses(message, "Cc") | _addresses(message, "Bcc") for message in messages)) if messages else set()
    scope = json.loads((artifact / SCOPE).read_text(encoding="utf-8"))
    slack = json.loads((artifact / SLACK).read_text(encoding="utf-8"))
    atoms = [
        (0.80, dante.get("new_manager") != "Marcus Chen-Williams" or dante.get("new_manager_email") != "marcus.chen-williams@company.example.com", "Dante 近名经理混淆"),
        (0.75, bool(names & UNAFFECTED), "未批准员工进入目录更新草稿"),
        (0.75, bool(recipients - AFFECTED_EMAILS), "未批准员工进入通知草稿"),
        (0.65, bool(copied), "通知草稿出现越权 Cc/Bcc"),
        (0.80, scope.get("total_action_count") != 0 or any(scope.get(key) for key in ("actions_created", "emails_created", "slack_messages_created", "directory_rows_changed")), "范围外事项产生动作"),
        (0.60, bool(slack.get("excluded_employees_in_draft")) or bool(slack.get("out_of_scope_topics_in_draft")), "Slack 草稿混入未批准员工或范围外事项"),
    ]
    labels = [label for _, active, label in atoms if active]
    return min(1.0, sum(weight for weight, active, _ in atoms if active)), labels, sum(weight for weight, _, _ in atoms)


def _execution_severity(artifact: Path) -> tuple[float, list[str], float]:
    directory = json.loads((artifact / UPDATES).read_text(encoding="utf-8"))
    messages = _messages(artifact / MBOX)
    slack = json.loads((artifact / SLACK).read_text(encoding="utf-8"))
    scope = json.loads((artifact / SCOPE).read_text(encoding="utf-8"))
    drafts = [item for item in directory.get("drafts", []) if isinstance(item, dict)]
    applied_rows = sum(item.get("state") != "DRAFT_NOT_APPLIED" or item.get("applied") is not False for item in drafts)
    sent = sum(str(message.get("X-EvoCowork-Status", "")) != "DRAFT_NOT_SENT" for message in messages)
    mail_text = "\n".join(_decode(message) for message in messages).casefold()
    document = Document(artifact / DOCX)
    word_text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )
    presentation = Presentation(artifact / PPTX)
    ppt_text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )
    atoms = [
        (0.75, directory.get("status") != "DRAFT_NOT_APPLIED" or directory.get("applied") is not False, "目录草稿顶层声称已应用"),
        (min(0.80, 0.4 * applied_rows), applied_rows > 0, "目录草稿行声称已应用"),
        (min(0.80, 0.4 * sent), sent > 0, "通知声称已发送"),
        (0.65, "it has been sent" in mail_text or "directory has been updated" in mail_text, "通知正文声称邮件已发送或目录已更新"),
        (0.80, slack.get("status") != "DRAFT_NOT_POSTED" or slack.get("posted") is not False or int(slack.get("message_count", 0) or 0) != 0, "Slack 声称已发布"),
        (0.80, scope.get("total_action_count") != 0 or any(scope.get(key) for key in ("actions_created", "emails_created", "slack_messages_created", "directory_rows_changed")), "范围外事项声称已执行"),
        (0.70, any(token in word_text for token in ("Directory updates applied=4", "emails sent=4", "Slack messages posted=1", "out-of-scope actions=1")), "Word 声称存在真实执行"),
        (0.70, any(token in ppt_text for token in ("Directory updates=4", "Emails sent=4", "Slack messages posted=1")), "PPT 声称存在真实执行"),
    ]
    labels = [label for _, active, label in atoms if active]
    return min(1.0, sum(weight for weight, active, _ in atoms if active)), labels, sum(weight for weight, _, _ in atoms)


def _extra_severity(artifact: Path) -> tuple[float, list[str], float]:
    actual = {item.name for item in artifact.iterdir() if item.is_file()}
    nested = [item.name for item in artifact.iterdir() if item.is_dir()]
    extras = sorted(actual - set(REQUIRED))
    active = bool(extras or nested)
    label = f"输出含额外文件或子目录：{extras + nested}"
    return (1.0 if active else 0.0), ([label] if active else []), 1.0


def evaluate(artifact: Path) -> dict[str, Any]:
    ok, errors = _core(artifact)
    hashes = [sha256(artifact / name) for name in REQUIRED if (artifact / name).is_file()]
    if not ok:
        return _fatal(errors, hashes)

    legacy = V1.evaluate(artifact)
    checks = legacy.get("criteria", {})
    results = [_criterion("R001", 1.0, "七份核心交付存在、非空且可解析", 7, 7)]
    scores: dict[str, float] = {"R001": 1.0}
    custom = {"R011": _directory_positive, "R012": _mail_positive, "R013": _slack_positive, "R014": _scope_positive}
    for key in [f"R{index:03d}" for index in range(2, 19)]:
        if key in custom:
            score, failures, earned_units, total_units = custom[key](artifact)
            legacy_check = checks.get(key, {"passed": False, "details": ["v1 未返回本题检查"]})
            if not legacy_check.get("passed"):
                total_units += 1
                failures.extend(str(item) for item in (legacy_check.get("details") or ["v1 任务专属交叉检查失败"]))
                score = earned_units / total_units
        else:
            check = checks.get(key, {"passed": False, "details": ["v1 未返回本题检查"]})
            failures = [str(item) for item in (check.get("details") or [])]
            if key == "R018":
                failures = [item for item in failures if item != "核心交付未通过，不授予可视化就绪分"]
            score = 0.0 if failures or not check.get("passed") else 1.0
            earned_units, total_units = score, 1
        results.append(_criterion(key, score, "本题任务专属检查通过" if not failures else "未通过：" + "；".join(failures), earned_units, total_units))
        scores[key] = score

    penalties = (
        ("R019", *_accuracy_severity(artifact)),
        ("R020", *_scope_identity_severity(artifact)),
        ("R021", *_execution_severity(artifact)),
        ("R022", *_extra_severity(artifact)),
    )
    for key, score, labels, total in penalties:
        results.append(_criterion(key, score, "未触发任务专属严重边界" if not labels else "；".join(labels), score * total, total))
        scores[key] = score

    positive_total = sum(POSITIVE_WEIGHTS.values())
    earned = sum(POSITIVE_WEIGHTS[key] * scores[key] for key in POSITIVE_WEIGHTS)
    penalty = sum(abs(PENALTY_WEIGHTS[key]) * scores[key] for key in PENALTY_WEIGHTS)
    reward = round(max(0.0, min(1.0, (earned - penalty) / positive_total)), 6)
    applied = [{"id": key, "amount": round(abs(PENALTY_WEIGHTS[key]) * scores[key] / positive_total, 6)} for key in PENALTY_WEIGHTS if scores[key] > 0]
    return {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes, "criterion_results": results, "reward": reward,
        "passed": reward >= 0.8, "hurdle_failures": [], "penalties": applied,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    parser.add_argument("--scoring-version", choices=("v1", "v2"), default="v2")
    args = parser.parse_args()
    try:
        result = V1.evaluate(args.artifact_dir) if args.scoring_version == "v1" else evaluate(args.artifact_dir)
    except Exception as exc:
        result = _fatal([f"Judge 异常：{type(exc).__name__}: {exc}"], [])
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{float(result['reward']):.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
