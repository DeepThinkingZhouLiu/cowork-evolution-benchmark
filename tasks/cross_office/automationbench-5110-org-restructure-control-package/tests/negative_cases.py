#!/usr/bin/env python3
"""向 AutomationBench 5110 Golden 注入本题专属错误，验证 Judge 拒绝坏答案。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
import zipfile
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import getaddresses
from io import BytesIO
from pathlib import Path
from typing import Callable

from docx import Document
from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
XLSX = "Org_Restructure_Audit.xlsx"
DOCX = "Org_Restructure_Control_Brief.docx"
PPTX = "Org_Restructure_Chart.pptx"
MBOX = "Org_Change_Notification_Drafts.mbox"
UPDATES = "Directory_Update_Drafts.json"
SLACK = "HR_Ops_Reorg_Draft.json"
SCOPE = "Scope_Control.json"
Mutator = Callable[[Path], None]
Case = tuple[str, Mutator, str]


def load_judge():
    spec = importlib.util.spec_from_file_location("reorg_5110_judge", ROOT / "tests" / "judge.py")
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入本题 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    module.render_office = lambda path: ({".xlsx": 8, ".docx": 3, ".pptx": 3}[path.suffix], "render-ok")
    return module


def rewrite_zip(path: Path, transform: Callable[[str, bytes], bytes | None]) -> None:
    temporary = path.with_suffix(path.suffix + ".negative")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = transform(info.filename, source.read(info.filename))
            if data is not None:
                target.writestr(info, data)
    temporary.replace(path)


def replace_ooxml_text(path: Path, old: str, new: str) -> None:
    old_bytes, new_bytes = old.encode("utf-8"), new.encode("utf-8")
    changed = False

    def transform(name: str, data: bytes) -> bytes:
        nonlocal changed
        if name.endswith(".xml") and old_bytes in data:
            changed = True
            return data.replace(old_bytes, new_bytes, 1)
        return data

    rewrite_zip(path, transform)
    if not changed:
        raise ValueError(f"OOXML 中找不到 {old!r}")


def sheet_xml_path(archive: zipfile.ZipFile, sheet_name: str) -> str:
    workbook = etree.fromstring(archive.read("xl/workbook.xml"))
    relationships = etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    ns = {
        "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    rel_id = workbook.xpath(f'//m:sheet[@name="{sheet_name}"]/@r:id', namespaces=ns)[0]
    target = relationships.xpath(f'//pr:Relationship[@Id="{rel_id}"]/@Target', namespaces=ns)[0]
    return target.lstrip("/") if target.startswith("/") else "xl/" + target.lstrip("/")


def edit_cell(path: Path, sheet_name: str, coordinate: str, *, value: str | int | bool | None = None, remove_formula: bool = False) -> None:
    with zipfile.ZipFile(path, "r") as archive:
        target_sheet = sheet_xml_path(archive, sheet_name)
    changed = False
    ns_uri = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

    def transform(name: str, data: bytes) -> bytes:
        nonlocal changed
        if name != target_sheet:
            return data
        root = etree.fromstring(data)
        cells = root.xpath(f'//m:c[@r="{coordinate}"]', namespaces={"m": ns_uri})
        if not cells:
            raise ValueError(f"{sheet_name}!{coordinate} 不存在")
        cell = cells[0]
        if remove_formula:
            formula = cell.find(f"{{{ns_uri}}}f")
            if formula is None:
                raise ValueError(f"{sheet_name}!{coordinate} 没有公式")
            cell.remove(formula)
        if value is not None:
            node = cell.find(f"{{{ns_uri}}}v")
            if node is None:
                node = etree.SubElement(cell, f"{{{ns_uri}}}v")
            if isinstance(value, bool):
                node.text = "1" if value else "0"
                cell.set("t", "b")
            elif isinstance(value, int):
                node.text = str(value)
                cell.attrib.pop("t", None)
            else:
                node.text = value
                cell.set("t", "str")
        changed = True
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    rewrite_zip(path, transform)
    if not changed:
        raise ValueError("目标单元格未修改")


def openpyxl_mutation(path: Path, mutator: Callable) -> None:
    workbook = load_workbook(path)
    mutator(workbook)
    workbook.save(path)


def mutate_json(path: Path, mutator: Callable[[dict], None]) -> None:
    payload = json.loads(path.read_text(encoding="utf-8"))
    mutator(payload)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def decode_message(message) -> str:
    payload = message.get_payload(decode=True) or b""
    return payload.decode(message.get_content_charset() or "utf-8", errors="replace")


def message_recipient(message) -> str:
    values = [address.casefold() for _, address in getaddresses(message.get_all("To", [])) if address]
    return values[0] if len(values) == 1 else ""


def mbox_states(path: Path) -> list[dict]:
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    return [{
        "from": str(message.get("From", "")), "to": str(message.get("To", "")),
        "cc": str(message.get("Cc", "")), "bcc": str(message.get("Bcc", "")),
        "subject": str(message.get("Subject", "")), "date": str(message.get("Date", "")),
        "message_id": str(message.get("Message-ID", "")), "status": str(message.get("X-EvoCowork-Status", "")),
        "body": decode_message(message), "recipient": message_recipient(message),
    } for message in messages]


def write_mbox_states(path: Path, states: list[dict]) -> None:
    chunks: list[bytes] = []
    for state in states:
        message = EmailMessage(policy=SMTP)
        for key, header in (("from", "From"), ("to", "To"), ("cc", "Cc"), ("bcc", "Bcc"), ("subject", "Subject"), ("date", "Date"), ("message_id", "Message-ID"), ("status", "X-EvoCowork-Status")):
            if state.get(key):
                message[header] = state[key]
        message.set_content(state["body"].rstrip() + "\n", charset="utf-8")
        buffer = BytesIO()
        BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
        chunks.extend([b"From MAILER-DAEMON Wed Apr 15 09:00:00 2026\n", buffer.getvalue().replace(b"\r\n", b"\n"), b"\n"])
    path.write_bytes(b"".join(chunks))


def mutate_mbox(path: Path, selector: str, mutator: Callable[[dict], None]) -> None:
    states = mbox_states(path)
    matching = [state for state in states if state["recipient"] == selector]
    if len(matching) != 1:
        raise ValueError(f"无法唯一找到 MBOX 收件人 {selector}")
    mutator(matching[0])
    write_mbox_states(path, states)


def drop_mbox(path: Path, selector: str) -> None:
    states = [state for state in mbox_states(path) if state["recipient"] != selector]
    if len(states) != 3:
        raise ValueError("未唯一删除通知草稿")
    write_mbox_states(path, states)


def body_replace(state: dict, old: str, new: str) -> None:
    if old not in state["body"]:
        raise ValueError(f"MBOX 正文中找不到 {old!r}")
    state["body"] = state["body"].replace(old, new, 1)


def replace_docx_text(path: Path, old: str, new: str, *, row_token: str | None = None) -> None:
    document = Document(path)
    paragraphs = list(document.paragraphs) if row_token is None else []
    for table in document.tables:
        for row in table.rows:
            if row_token is not None and row_token not in " | ".join(cell.text for cell in row.cells):
                continue
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        if old in paragraph.text:
            paragraph.text = paragraph.text.replace(old, new, 1)
            document.save(path)
            return
    raise ValueError(f"Word 中找不到 {old!r}")


def replace_pptx_text(path: Path, old: str, new: str) -> None:
    presentation = Presentation(path)
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and old in shape.text:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if old in run.text:
                            run.text = run.text.replace(old, new, 1)
                            presentation.save(path)
                            return
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        if old in cell.text:
                            cell.text = cell.text.replace(old, new, 1)
                            presentation.save(path)
                            return
    raise ValueError(f"PPT 中找不到 {old!r}")


def drop_ppt_slide(path: Path) -> None:
    presentation = Presentation(path)
    slide_id = presentation.slides._sldIdLst[-1]
    presentation.part.drop_rel(slide_id.rId)
    presentation.slides._sldIdLst.remove(slide_id)
    presentation.save(path)


def remove_ppt_connector(path: Path) -> None:
    presentation = Presentation(path)
    slide = presentation.slides[1]
    for shape in list(slide.shapes):
        if shape.shape_type == 9:
            slide.shapes._spTree.remove(shape._element)
            presentation.save(path)
            return
    raise ValueError("PPT 中找不到连接线")


def build_cases() -> list[Case]:
    cases: list[Case] = [
        ("missing-xlsx", lambda root: (root / XLSX).unlink(), "R001"),
        ("missing-docx", lambda root: (root / DOCX).unlink(), "R001"),
        ("missing-pptx", lambda root: (root / PPTX).unlink(), "R001"),
        ("missing-mbox", lambda root: (root / MBOX).unlink(), "R001"),
        ("missing-directory-json", lambda root: (root / UPDATES).unlink(), "R001"),
        ("missing-slack-json", lambda root: (root / SLACK).unlink(), "R001"),
        ("missing-scope-json", lambda root: (root / SCOPE).unlink(), "R001"),
        ("extra-debug-file", lambda root: (root / "debug.txt").write_text("debug\n", encoding="utf-8"), "R001"),
        ("corrupt-xlsx", lambda root: (root / XLSX).write_bytes(b"not-an-xlsx"), "R001"),
        ("corrupt-docx", lambda root: (root / DOCX).write_bytes(b"not-a-docx"), "R001"),
        ("corrupt-pptx", lambda root: (root / PPTX).write_bytes(b"not-a-pptx"), "R001"),
        ("source-aisha-name", lambda root: replace_ooxml_text(root / XLSX, "Aisha Jordan", "Aisha Jordon"), "R002"),
        ("source-bo-email", lambda root: replace_ooxml_text(root / XLSX, "bo.lindqvist@company.example.com", "bo.wrong@company.example.com"), "R002"),
        ("source-chiara-department", lambda root: edit_cell(root / XLSX, "Source_Directory", "D4", value="Product"), "R002"),
        ("source-dante-manager", lambda root: edit_cell(root / XLSX, "Source_Directory", "E5", value="Marcus Chen"), "R002"),
        ("source-elena-manager-email", lambda root: replace_ooxml_text(root / XLSX, "elena.vasquez@company.example.com", "elena.bad@company.example.com"), "R002"),
        ("source-felix-manager", lambda root: edit_cell(root / XLSX, "Source_Directory", "E7", value="Lisa Wang"), "R002"),
        ("source-grace-near-name", lambda root: edit_cell(root / XLSX, "Source_Directory", "E8", value="Marcus Chen"), "R002"),
        ("source-hugo-row", lambda root: edit_cell(root / XLSX, "Source_Directory", "A9", value=10), "R002"),
        ("approved-order", lambda root: edit_cell(root / XLSX, "Approved_Changes", "A2", value=2), "R003"),
        ("approved-effective-date", lambda root: replace_ooxml_text(root / XLSX, "2026-04-14", "2026-04-15"), "R003"),
        ("approved-aisha-old", lambda root: edit_cell(root / XLSX, "Approved_Changes", "D2", value="Derek Huang"), "R003"),
        ("approved-aisha-new", lambda root: edit_cell(root / XLSX, "Approved_Changes", "E2", value="Lisa Wang"), "R003"),
        ("approved-bo-new-email", lambda root: edit_cell(root / XLSX, "Approved_Changes", "F3", value="derek.wrong@company.example.com"), "R003"),
        ("approved-chiara-new", lambda root: edit_cell(root / XLSX, "Approved_Changes", "E4", value="Marcus Chen"), "R003"),
        ("approved-dante-near-name", lambda root: edit_cell(root / XLSX, "Approved_Changes", "E5", value="Marcus Chen"), "R008"),
        ("approved-dante-email", lambda root: edit_cell(root / XLSX, "Approved_Changes", "F5", value="marcus.chen@company.example.com"), "R008"),
        ("approved-trace", lambda root: edit_cell(root / XLSX, "Approved_Changes", "G5", value="unknown"), "R003"),
        ("manager-lisa-email", lambda root: edit_cell(root / XLSX, "Manager_Lookup", "B2", value="lisa.wrong@company.example.com"), "R009"),
        ("manager-marcus-email", lambda root: edit_cell(root / XLSX, "Manager_Lookup", "B3", value="marcus.chen-williams@company.example.com"), "R009"),
        ("manager-marcus-control", lambda root: edit_cell(root / XLSX, "Manager_Lookup", "E3", value="SAME_PERSON"), "R009"),
        ("manager-derek-rows", lambda root: edit_cell(root / XLSX, "Manager_Lookup", "C4", value="4"), "R009"),
        ("manager-williams-name", lambda root: edit_cell(root / XLSX, "Manager_Lookup", "A5", value="Marcus Chen"), "R009"),
        ("manager-williams-email", lambda root: edit_cell(root / XLSX, "Manager_Lookup", "B5", value="marcus.chen@company.example.com"), "R009"),
        ("evidence-email-truncated", lambda root: replace_ooxml_text(root / XLSX, "Marcus Chen (Engineering) and Marcus Chen-Williams (Product) are different managers", "Both Marcus names are the same manager"), "R003"),
        ("evidence-time", lambda root: replace_ooxml_text(root / XLSX, "2026-04-15T09:00:00Z", "2026-04-16T09:00:00Z"), "R003"),
        ("evidence-channel", lambda root: replace_ooxml_text(root / XLSX, "C_HROPS_5110", "C_WRONG_5110"), "R003"),
        ("evidence-source-trace", lambda root: edit_cell(root / XLSX, "Evidence_Register", "E2", value="unknown"), "R017"),
    ]

    formula_cells = [
        ("B2", "source-name"), ("F2", "source-manager-email"), ("G2", "affected"),
        ("H2", "approved-old"), ("I2", "old-match"), ("J2", "approved-new"),
        ("K2", "new-email"), ("L2", "updated-manager"), ("M2", "updated-email"),
        ("N2", "decision"), ("O2", "update-route"), ("P2", "notification-route"),
        ("Q2", "slack-route"), ("R2", "trace"), ("L6", "unaffected-fallback"),
        ("M6", "unaffected-email-fallback"),
    ]
    cases.extend((f"formula-{label}-staticized", lambda root, cell=cell: edit_cell(root / XLSX, "Reorg_Decisions", cell, remove_formula=True), "R004") for cell, label in formula_cells)
    queue_formula_cells = [
        ("Directory_Update_Queue", "A2", "directory-id"), ("Directory_Update_Queue", "H2", "directory-manager"),
        ("Directory_Update_Queue", "K2", "directory-state"), ("Directory_Update_Queue", "L2", "directory-applied"),
        ("Notification_Queue", "A2", "notification-id"), ("Notification_Queue", "E2", "notification-recipient"),
        ("Notification_Queue", "H2", "notification-manager"), ("Notification_Queue", "K2", "notification-state"),
        ("Notification_Queue", "L2", "notification-sent"), ("Control_Summary", "B5", "summary-conflicts"),
        ("Control_Summary", "B9", "summary-applied"), ("Control_Summary", "B12", "summary-scope"),
    ]
    cases.extend((f"formula-{label}-staticized", lambda root, sheet=sheet, cell=cell: edit_cell(root / XLSX, sheet, cell, remove_formula=True), "R004") for sheet, cell, label in queue_formula_cells)

    cases.extend([
        ("decision-aisha-manager", lambda root: edit_cell(root / XLSX, "Reorg_Decisions", "L2", value="Lisa Wang"), "R005"),
        ("decision-aisha-email", lambda root: edit_cell(root / XLSX, "Reorg_Decisions", "M2", value="lisa.wang@company.example.com"), "R005"),
        ("decision-bo-manager", lambda root: edit_cell(root / XLSX, "Reorg_Decisions", "L3", value="Marcus Chen"), "R006"),
        ("decision-bo-route", lambda root: edit_cell(root / XLSX, "Reorg_Decisions", "P3", value="NO"), "R006"),
        ("decision-chiara-manager", lambda root: edit_cell(root / XLSX, "Reorg_Decisions", "L4", value="Derek Huang"), "R007"),
        ("decision-chiara-result", lambda root: edit_cell(root / XLSX, "Reorg_Decisions", "N4", value="SOURCE_CONFLICT"), "R007"),
        ("decision-dante-near-name", lambda root: edit_cell(root / XLSX, "Reorg_Decisions", "L5", value="Marcus Chen"), "R008"),
        ("decision-dante-near-email", lambda root: edit_cell(root / XLSX, "Reorg_Decisions", "M5", value="marcus.chen@company.example.com"), "R008"),
        ("decision-elena-affected", lambda root: edit_cell(root / XLSX, "Reorg_Decisions", "G6", value="YES"), "R010"),
        ("decision-elena-manager", lambda root: edit_cell(root / XLSX, "Reorg_Decisions", "L6", value="Lisa Wang"), "R010"),
        ("decision-felix-update", lambda root: edit_cell(root / XLSX, "Reorg_Decisions", "O7", value="YES"), "R010"),
        ("decision-grace-manager", lambda root: edit_cell(root / XLSX, "Reorg_Decisions", "L8", value="Marcus Chen"), "R010"),
        ("decision-hugo-notify", lambda root: edit_cell(root / XLSX, "Reorg_Decisions", "P9", value="YES"), "R010"),
        ("summary-approved", lambda root: edit_cell(root / XLSX, "Control_Summary", "B3", value=5), "R004"),
        ("summary-unchanged", lambda root: edit_cell(root / XLSX, "Control_Summary", "B4", value=3), "R004"),
        ("summary-applied", lambda root: edit_cell(root / XLSX, "Control_Summary", "B9", value=1), "R004"),
        ("summary-sent", lambda root: edit_cell(root / XLSX, "Control_Summary", "B10", value=1), "R004"),
        ("summary-posted", lambda root: edit_cell(root / XLSX, "Control_Summary", "B11", value=1), "R004"),
        ("summary-scope-actions", lambda root: edit_cell(root / XLSX, "Control_Summary", "B12", value=1), "R004"),
        ("excel-directory-applied", lambda root: edit_cell(root / XLSX, "Directory_Update_Queue", "L2", value=True), "R011"),
        ("excel-directory-state", lambda root: edit_cell(root / XLSX, "Directory_Update_Queue", "K3", value="APPLIED"), "R011"),
        ("excel-directory-near-name", lambda root: edit_cell(root / XLSX, "Directory_Update_Queue", "H5", value="Marcus Chen"), "R008"),
        ("excel-notification-sent", lambda root: edit_cell(root / XLSX, "Notification_Queue", "L2", value=True), "R012"),
        ("excel-notification-recipient", lambda root: edit_cell(root / XLSX, "Notification_Queue", "E4", value="hugo.fernandez@company.example.com"), "R012"),
        ("native-table-removed", lambda root: openpyxl_mutation(root / XLSX, lambda wb: wb["Source_Directory"].tables.clear()), "R018"),
        ("native-chart-removed", lambda root: openpyxl_mutation(root / XLSX, lambda wb: setattr(wb["Control_Summary"], "_charts", [])), "R018"),
        ("conditional-formatting-removed", lambda root: openpyxl_mutation(root / XLSX, lambda wb: setattr(wb["Reorg_Decisions"], "conditional_formatting", type(wb["Reorg_Decisions"].conditional_formatting)())), "R018"),
    ])
    cases.extend([
        ("directory-top-applied", lambda root: mutate_json(root / UPDATES, lambda p: p.update({"status": "APPLIED", "applied": True})), "R011"),
        ("directory-processing-time", lambda root: mutate_json(root / UPDATES, lambda p: p.update({"processing_time": "2026-04-16T09:00:00Z"})), "R011"),
        ("directory-effective-date", lambda root: mutate_json(root / UPDATES, lambda p: p.update({"effective_date": "2026-04-15"})), "R011"),
        ("directory-count", lambda root: mutate_json(root / UPDATES, lambda p: p.update({"draft_count": 5})), "R011"),
        ("directory-drop-aisha", lambda root: mutate_json(root / UPDATES, lambda p: p["drafts"].pop(0)), "R011"),
        ("directory-duplicate-bo", lambda root: mutate_json(root / UPDATES, lambda p: p["drafts"].append(dict(p["drafts"][1]))), "R011"),
        ("directory-aisha-manager", lambda root: mutate_json(root / UPDATES, lambda p: p["drafts"][0].update({"new_manager": "Lisa Wang"})), "R005"),
        ("directory-bo-email", lambda root: mutate_json(root / UPDATES, lambda p: p["drafts"][1].update({"new_manager_email": "derek.wrong@company.example.com"})), "R006"),
        ("directory-chiara-old", lambda root: mutate_json(root / UPDATES, lambda p: p["drafts"][2].update({"old_manager": "Lisa Wang"})), "R007"),
        ("directory-dante-near-name", lambda root: mutate_json(root / UPDATES, lambda p: p["drafts"][3].update({"new_manager": "Marcus Chen", "new_manager_email": "marcus.chen@company.example.com"})), "R008"),
        ("directory-row-applied", lambda root: mutate_json(root / UPDATES, lambda p: p["drafts"][0].update({"state": "APPLIED", "applied": True})), "R011"),
        ("directory-trace-missing", lambda root: mutate_json(root / UPDATES, lambda p: p["drafts"][2].update({"source_trace": []})), "R007"),
        ("directory-add-elena", lambda root: mutate_json(root / UPDATES, lambda p: p["drafts"].append({**p["drafts"][0], "employee": "Elena Vasquez", "employee_email": "elena.vasquez@company.example.com", "source_row": 6})), "R010"),
    ])

    cases.extend([
        ("mbox-drop-aisha", lambda root: drop_mbox(root / MBOX, "aisha.jordan@company.example.com"), "R005"),
        ("mbox-drop-dante", lambda root: drop_mbox(root / MBOX, "dante.obi@company.example.com"), "R008"),
        ("mbox-aisha-recipient", lambda root: mutate_mbox(root / MBOX, "aisha.jordan@company.example.com", lambda s: s.update({"to": "elena.vasquez@company.example.com"})), "R012"),
        ("mbox-bo-cc", lambda root: mutate_mbox(root / MBOX, "bo.lindqvist@company.example.com", lambda s: s.update({"cc": "hr-all@company.example.com"})), "R012"),
        ("mbox-chiara-bcc", lambda root: mutate_mbox(root / MBOX, "chiara.russo@company.example.com", lambda s: s.update({"bcc": "vp-people@company.example.com"})), "R012"),
        ("mbox-aisha-status", lambda root: mutate_mbox(root / MBOX, "aisha.jordan@company.example.com", lambda s: s.update({"status": "SENT"})), "R012"),
        ("mbox-aisha-department", lambda root: mutate_mbox(root / MBOX, "aisha.jordan@company.example.com", lambda s: body_replace(s, "Department=Engineering", "Department=Product")), "R005"),
        ("mbox-aisha-old-manager", lambda root: mutate_mbox(root / MBOX, "aisha.jordan@company.example.com", lambda s: body_replace(s, "Old Manager=Lisa Wang", "Old Manager=Derek Huang")), "R005"),
        ("mbox-bo-new-manager", lambda root: mutate_mbox(root / MBOX, "bo.lindqvist@company.example.com", lambda s: body_replace(s, "New Manager=Derek Huang", "New Manager=Marcus Chen")), "R006"),
        ("mbox-bo-new-email", lambda root: mutate_mbox(root / MBOX, "bo.lindqvist@company.example.com", lambda s: body_replace(s, "New Manager Email=derek.huang@company.example.com", "New Manager Email=derek.wrong@company.example.com")), "R006"),
        ("mbox-chiara-effective-date", lambda root: mutate_mbox(root / MBOX, "chiara.russo@company.example.com", lambda s: body_replace(s, "Effective Date=2026-04-14", "Effective Date=2026-04-15")), "R007"),
        ("mbox-chiara-source-row", lambda root: mutate_mbox(root / MBOX, "chiara.russo@company.example.com", lambda s: body_replace(s, "Source Row=4", "Source Row=5")), "R007"),
        ("mbox-dante-near-name", lambda root: mutate_mbox(root / MBOX, "dante.obi@company.example.com", lambda s: body_replace(s, "New Manager=Marcus Chen-Williams", "New Manager=Marcus Chen")), "R008"),
        ("mbox-dante-near-email", lambda root: mutate_mbox(root / MBOX, "dante.obi@company.example.com", lambda s: body_replace(s, "New Manager Email=marcus.chen-williams@company.example.com", "New Manager Email=marcus.chen@company.example.com")), "R008"),
        ("mbox-false-sent-claim", lambda root: mutate_mbox(root / MBOX, "aisha.jordan@company.example.com", lambda s: body_replace(s, "has not been sent", "has been sent")), "R012"),
        ("mbox-false-directory-claim", lambda root: mutate_mbox(root / MBOX, "bo.lindqvist@company.example.com", lambda s: body_replace(s, "directory has not been updated", "directory has been updated")), "R012"),
        ("mbox-add-unaffected-name", lambda root: mutate_mbox(root / MBOX, "chiara.russo@company.example.com", lambda s: s.update({"body": s["body"] + "\nHugo Fernandez\n"})), "R010"),
    ])

    cases.extend([
        ("slack-wrong-channel", lambda root: mutate_json(root / SLACK, lambda p: p["source_channel"].update({"id": "C_WRONG"})), "R013"),
        ("slack-posted", lambda root: mutate_json(root / SLACK, lambda p: p.update({"status": "POSTED", "posted": True, "message_count": 1})), "R013"),
        ("slack-draft-count", lambda root: mutate_json(root / SLACK, lambda p: p.update({"draft_count": 2})), "R013"),
        ("slack-change-count", lambda root: mutate_json(root / SLACK, lambda p: p["draft_message"].update({"change_count": 3})), "R013"),
        ("slack-drop-aisha-line", lambda root: mutate_json(root / SLACK, lambda p: p["draft_message"]["change_lines"].pop(0)), "R013"),
        ("slack-reverse-order", lambda root: mutate_json(root / SLACK, lambda p: p["draft_message"]["change_lines"].reverse()), "R013"),
        ("slack-aisha-manager", lambda root: mutate_json(root / SLACK, lambda p: p["draft_message"]["change_lines"].__setitem__(0, p["draft_message"]["change_lines"][0].replace("Marcus Chen", "Lisa Wang"))), "R013"),
        ("slack-dante-near-name", lambda root: mutate_json(root / SLACK, lambda p: p["draft_message"]["change_lines"].__setitem__(3, p["draft_message"]["change_lines"][3].replace("Marcus Chen-Williams", "Marcus Chen"))), "R013"),
        ("slack-add-hugo", lambda root: mutate_json(root / SLACK, lambda p: p["draft_message"]["change_lines"].append("5. Hugo Fernandez: Lisa Wang -> Derek Huang")), "R013"),
        ("slack-excluded-employee", lambda root: mutate_json(root / SLACK, lambda p: p.update({"excluded_employees_in_draft": ["Elena Vasquez"]})), "R010"),
        ("slack-add-office-move", lambda root: mutate_json(root / SLACK, lambda p: p.update({"out_of_scope_topics_in_draft": ["office_move"]})), "R014"),
        ("slack-false-execution", lambda root: mutate_json(root / SLACK, lambda p: p["draft_message"].update({"execution_note": "POSTED; directory updates applied=4; employee emails sent=4"})), "R013"),
        ("slack-trace", lambda root: mutate_json(root / SLACK, lambda p: p.update({"source_trace": []})), "R017"),
    ])

    cases.extend([
        ("scope-topic-count", lambda root: mutate_json(root / SCOPE, lambda p: p.update({"total_out_of_scope_topics": 2})), "R014"),
        ("scope-action-count", lambda root: mutate_json(root / SCOPE, lambda p: p.update({"total_action_count": 1})), "R014"),
        ("scope-drop-product-launch", lambda root: mutate_json(root / SCOPE, lambda p: p["topics"].pop(0)), "R014"),
        ("scope-product-in-scope", lambda root: mutate_json(root / SCOPE, lambda p: p["topics"][0].update({"decision": "IN_SCOPE", "action_count": 1})), "R014"),
        ("scope-office-owner", lambda root: mutate_json(root / SCOPE, lambda p: p["topics"][1].update({"owner": "HR Ops"})), "R014"),
        ("scope-sprint-owner", lambda root: mutate_json(root / SCOPE, lambda p: p["topics"][2].update({"owner": "Engineering Manager"})), "R014"),
        ("scope-product-excerpt", lambda root: mutate_json(root / SCOPE, lambda p: p["topics"][0].update({"prompt_excerpt": "Product launch"})), "R014"),
        ("scope-created-action", lambda root: mutate_json(root / SCOPE, lambda p: p["actions_created"].append({"topic": "office_move"})), "R014"),
        ("scope-created-email", lambda root: mutate_json(root / SCOPE, lambda p: p["emails_created"].append("marketing@company.example.com")), "R014"),
        ("scope-created-slack", lambda root: mutate_json(root / SCOPE, lambda p: p["slack_messages_created"].append("launch")), "R014"),
        ("scope-changed-directory", lambda root: mutate_json(root / SCOPE, lambda p: p["directory_rows_changed"].append(6)), "R014"),
        ("scope-trace", lambda root: mutate_json(root / SCOPE, lambda p: p.update({"source_trace": []})), "R017"),
    ])

    cases.extend([
        ("word-aisha-manager", lambda root: replace_docx_text(root / DOCX, "Marcus Chen", "Lisa Wang", row_token="Aisha Jordan"), "R005"),
        ("word-bo-manager", lambda root: replace_docx_text(root / DOCX, "Derek Huang", "Marcus Chen", row_token="Bo Lindqvist"), "R006"),
        ("word-chiara-manager", lambda root: replace_docx_text(root / DOCX, "Lisa Wang", "Derek Huang", row_token="Chiara Russo"), "R007"),
        ("word-dante-near-name", lambda root: replace_docx_text(root / DOCX, "Marcus Chen-Williams", "Marcus Chen", row_token="Dante Obi"), "R008"),
        ("word-elena-manager", lambda root: replace_docx_text(root / DOCX, "Derek Huang", "Lisa Wang", row_token="Elena Vasquez"), "R010"),
        ("word-manager-group", lambda root: replace_docx_text(root / DOCX, "Dante Obi; Grace Kim", "Dante Obi; Felix Braun", row_token="Marcus Chen-Williams"), "R009"),
        ("word-effective-date", lambda root: replace_docx_text(root / DOCX, "approved 2026-04-14", "approved 2026-04-15"), "R016"),
        ("word-directory-applied", lambda root: replace_docx_text(root / DOCX, "Directory updates applied=0", "Directory updates applied=4"), "R016"),
        ("word-emails-sent", lambda root: replace_docx_text(root / DOCX, "emails sent=0", "emails sent=4"), "R016"),
        ("word-slack-posted", lambda root: replace_docx_text(root / DOCX, "Slack messages posted=0", "Slack messages posted=1"), "R016"),
        ("word-scope-action", lambda root: replace_docx_text(root / DOCX, "out-of-scope actions=0", "out-of-scope actions=1"), "R016"),
    ])

    cases.extend([
        ("ppt-drop-slide", lambda root: drop_ppt_slide(root / PPTX), "R015"),
        ("ppt-drop-connector", lambda root: remove_ppt_connector(root / PPTX), "R015"),
        ("ppt-aisha-manager", lambda root: replace_pptx_text(root / PPTX, "Marcus Chen", "Lisa Wang"), "R015"),
        ("ppt-bo-manager", lambda root: replace_pptx_text(root / PPTX, "Derek Huang", "Marcus Chen"), "R015"),
        ("ppt-dante-near-name", lambda root: replace_pptx_text(root / PPTX, "Marcus Chen-Williams", "Marcus Chen"), "R015"),
        ("ppt-hugo-preserved", lambda root: replace_pptx_text(root / PPTX, "Hugo Fernandez", "Hugo Wrong"), "R010"),
        ("ppt-manager-email", lambda root: replace_pptx_text(root / PPTX, "marcus.chen-williams@company.example.com", "marcus.chen@company.example.com"), "R015"),
        ("ppt-effective-date", lambda root: replace_pptx_text(root / PPTX, "2026-04-14", "2026-04-15"), "R015"),
        ("ppt-scope-decision", lambda root: replace_pptx_text(root / PPTX, "OUT_OF_SCOPE", "IN_SCOPE"), "R014"),
        ("ppt-directory-executed", lambda root: replace_pptx_text(root / PPTX, "Directory updates=0", "Directory updates=4"), "R015"),
        ("ppt-emails-sent", lambda root: replace_pptx_text(root / PPTX, "Emails sent=0", "Emails sent=4"), "R015"),
        ("ppt-slack-posted", lambda root: replace_pptx_text(root / PPTX, "Slack messages posted=0", "Slack messages posted=1"), "R015"),
    ])
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--json-out", type=Path)
    args = parser.parse_args()
    judge = load_judge()
    results: list[dict] = []
    failures: list[str] = []
    cases = build_cases()
    for name, mutator, expected_failure in cases:
        with tempfile.TemporaryDirectory(prefix=f"negative-5110-{name}-") as temporary:
            root = Path(temporary) / "artifacts"
            shutil.copytree(GOLDEN, root)
            try:
                mutator(root)
                result = judge.evaluate(root)
                detected = not result.get("criteria", {}).get(expected_failure, {}).get("passed", True)
                rejected = not result.get("passed", False) and float(result.get("reward", 1.0)) < 0.8
                results.append({
                    "case": name, "expected_failure": expected_failure,
                    "detected": detected, "rejected": rejected,
                    "reward": result.get("reward"), "hurdle_failures": result.get("hurdle_failures", []),
                })
                if not detected or not rejected:
                    failures.append(f"{name}: expected={expected_failure}, detected={detected}, reward={result.get('reward')}")
            except Exception as exc:
                failures.append(f"{name}: mutation/evaluation error: {type(exc).__name__}: {exc}")
    payload = {
        "task_id": "cowork-evo/automationbench-5110-org-restructure-control-package",
        "case_count": len(results), "all_rejected": not failures and len(results) == len(cases),
        "max_reward": max((float(item["reward"]) for item in results), default=None),
        "failures": failures, "results": results,
    }
    if args.json_out:
        args.json_out.parent.mkdir(parents=True, exist_ok=True)
        args.json_out.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    if failures or len(results) != len(cases):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
