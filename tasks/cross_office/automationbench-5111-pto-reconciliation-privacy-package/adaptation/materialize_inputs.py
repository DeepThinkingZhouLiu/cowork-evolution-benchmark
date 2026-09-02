#!/usr/bin/env python3
"""只把 AutomationBench 5111 的公开初始状态静态化为本题输入。

本脚本只接受 ``hr.pto_reconciliation``，逐项断言八条 PTO 余额、
三条政策、两封完整邮件、冻结时间、两个 Slack 频道和十一条上游断言。
它不是通用制题器，不会创造员工、余额、邮箱、政策或修正事实。
"""

from __future__ import annotations

import hashlib
import importlib.util
import json
import shutil
import sys
import types
import zipfile
from datetime import datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
UPSTREAM = Path(__file__).with_name("upstream-hr-tasks.py")
SNAPSHOT = Path(__file__).with_name("upstream-task-5111.json")
LICENSE = Path(__file__).with_name("LICENSE-AutomationBench-MIT")
DATA_DIR = ROOT / "data" / "input_files"
ENV_DIR = ROOT / "environment" / "input_files"
FIXED_ZIP_TIME = (2026, 4, 15, 9, 0, 0)
SOURCE_SHA256 = "c173a36a03f43096bd122577584aa32f9df379d8760aa46c43be04c85d101779"
LICENSE_SHA256 = "c6f9d87aff6d627d5cd78cb09364dafed7d6f1c4bb7a91183d57377157e96ba4"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def freeze_upstream_assets() -> None:
    if not UPSTREAM.is_file():
        source = ROOT.parent / "automationbench-5110-org-restructure-control-package" / "adaptation" / "upstream-hr-tasks.py"
        if not source.is_file() or sha256(source) != SOURCE_SHA256:
            raise FileNotFoundError("找不到哈希匹配的 AutomationBench HR 冻结源文件")
        shutil.copy2(source, UPSTREAM)
    if sha256(UPSTREAM) != SOURCE_SHA256:
        raise ValueError("AutomationBench HR 冻结源文件哈希错误")
    if not LICENSE.is_file():
        source = ROOT.parent / "automationbench-5110-org-restructure-control-package" / "adaptation" / "LICENSE-AutomationBench-MIT"
        if not source.is_file() or sha256(source) != LICENSE_SHA256:
            raise FileNotFoundError("找不到哈希匹配的 AutomationBench LICENSE")
        shutil.copy2(source, LICENSE)
    if sha256(LICENSE) != LICENSE_SHA256:
        raise ValueError("AutomationBench LICENSE 哈希错误")


def load_task() -> dict:
    datasets = types.ModuleType("datasets")
    datasets.Dataset = object
    sys.modules["datasets"] = datasets
    for name in ("automationbench", "automationbench.domains", "automationbench.domains.hr"):
        sys.modules[name] = types.ModuleType(name)
    noise = types.ModuleType("automationbench.domains.hr._noise")
    noise.apply_noise = lambda value: value
    sys.modules[noise.__name__] = noise
    spec = importlib.util.spec_from_file_location("automationbench_hr_tasks_frozen_5111", UPSTREAM)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入冻结的 AutomationBench HR 源文件")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    task = module.get_task_5111()
    if task.get("example_id") != 5111 or task.get("task") != "hr.pto_reconciliation":
        raise ValueError("冻结函数不再是 AutomationBench 5111")
    return task


def assert_shape(task: dict) -> None:
    prompt = task["prompt"][1]["content"]
    assert all(token in prompt for token in ("Reconcile PTO balances for Q2", "Notify employees", "#company-announcements", "most recent data", "current policies"))
    state = task["info"]["initial_state"]
    assert state["meta"] == {"schema_version": "0.1.0", "current_time": "2026-04-15T09:00:00Z"}
    messages = state["gmail"]["messages"]
    assert [message["id"] for message in messages] == ["msg_pto_corrections_5111", "msg_pto_privacy_5111"]
    assert all(token in messages[0]["body_plain"] for token in ("Rita Johansson: Should be 3 days", "Sam Okafor: Should be -2 days", "update the tracker"))
    assert all(token in messages[1]["body_plain"] for token in ("individual PTO balances are confidential", "Do NOT post names or balances", "aggregate numbers only", "private #hr-general", "notified individually by email"))
    spreadsheet = state["google_sheets"]["spreadsheets"][0]
    assert spreadsheet["id"] == "ss_pto_5111" and [sheet["title"] for sheet in spreadsheet["worksheets"]] == ["Q2 Balances", "PTO Policy"]
    balances = spreadsheet["worksheets"][0]["rows"]
    assert [row["row_id"] for row in balances] == list(range(2, 10))
    assert [row["cells"]["Employee"] for row in balances] == ["Quinn Torres", "Rita Johansson", "Sam Okafor", "Tina Marchetti", "Uma Krishnan", "Victor Pham", "Wendy Liu", "Xavier Dunn"]
    assert [row["cells"]["PTO Balance (Days)"] for row in balances] == ["12", "8", "3", "15", "0", "7", "-1", "9"]
    assert balances[6]["cells"]["Notes"] == "Approved overdraft - manager signed off"
    policies = spreadsheet["worksheets"][1]["rows"]
    assert [row["cells"]["Rule"] for row in policies] == ["Negative Balance", "Balance Correction", "No Change Needed"]
    assert state["slack"] == {"channels": [{"id": "C_HR_5111", "name": "hr-general", "is_private": False}, {"id": "C_ANN_5111", "name": "company-announcements", "is_private": False}], "messages": []}
    assert len(task["info"]["assertions"]) == 11


def normalize_xlsx(path: Path) -> None:
    temporary = path.with_suffix(".normalized.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    temporary.replace(path)


def style_sheet(sheet, widths: dict[str, float]) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = sheet.dimensions


def write_tracker(state: dict, target: Path) -> None:
    spreadsheet = state["google_sheets"]["spreadsheets"][0]
    workbook = Workbook()
    workbook.creator = "AutomationBench 5111 public fixture"
    workbook.title = spreadsheet["title"]
    workbook.subject = "Static Office representation of PTO balances and policy"
    workbook.description = "No employee, balance, email, note or policy was added beyond AutomationBench 5111."
    workbook.remove(workbook.active)
    balances = spreadsheet["worksheets"][0]
    sheet = workbook.create_sheet(balances["title"])
    headers = ["Source Row", "Employee", "Employee Email", "Manager Email", "PTO Balance (Days)", "Status", "Notes"]
    sheet.append(headers)
    for item in balances["rows"]:
        cells = item["cells"]
        sheet.append([item["row_id"], cells["Employee"], cells["Employee Email"], cells["Manager Email"], int(cells["PTO Balance (Days)"]), cells["Status"], cells["Notes"]])
    table = Table(displayName="PTOBalanceSourceInput", ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    style_sheet(sheet, {"A": 12, "B": 23, "C": 38, "D": 36, "E": 20, "F": 16, "G": 46})
    policies = spreadsheet["worksheets"][1]
    policy_sheet = workbook.create_sheet(policies["title"])
    policy_sheet.append(["Source Row", "Rule", "Action"])
    for item in policies["rows"]:
        policy_sheet.append([item["row_id"], item["cells"]["Rule"], item["cells"]["Action"]])
    table = Table(displayName="PTOPolicySourceInput", ref=policy_sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    policy_sheet.add_table(table)
    style_sheet(policy_sheet, {"A": 12, "B": 25, "C": 125})
    workbook.save(target)
    normalize_xlsx(target)


def message_bytes(record: dict) -> bytes:
    message = EmailMessage(policy=SMTP)
    message["From"] = record["from_"]
    message["To"] = ", ".join(record["to"])
    message["Subject"] = record["subject"]
    parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
    message["Date"] = parsed.strftime("%a, %d %b %Y %H:%M:%S +0000")
    message["Message-ID"] = f"<{record['id']}@automationbench.invalid>"
    message["X-AutomationBench-Message-ID"] = record["id"]
    message["X-AutomationBench-Thread-ID"] = record["thread_id"]
    message["X-AutomationBench-Labels"] = ",".join(record.get("label_ids", []))
    message.set_content(record["body_plain"], charset="utf-8")
    buffer = BytesIO()
    BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
    return buffer.getvalue()


def write_mbox(messages: list[dict], target: Path) -> None:
    chunks: list[bytes] = []
    for record in messages:
        parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
        chunks.append(f"From {record['from_']} {parsed.strftime('%a %b %d %H:%M:%S %Y')}\n".encode("ascii"))
        chunks.append(message_bytes(record).replace(b"\r\n", b"\n"))
        chunks.append(b"\n")
    target.write_bytes(b"".join(chunks))


def main() -> None:
    freeze_upstream_assets()
    task = load_task()
    assert_shape(task)
    state = task["info"]["initial_state"]
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    ENV_DIR.mkdir(parents=True, exist_ok=True)
    SNAPSHOT.write_text(json.dumps(task, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_tracker(state, DATA_DIR / "PTO_Balance_Tracker.xlsx")
    write_mbox(state["gmail"]["messages"], DATA_DIR / "PTO_Corrections_and_Privacy.mbox")
    (DATA_DIR / "Slack_Channels.json").write_text(json.dumps(state["slack"], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (DATA_DIR / "Workflow_Context.json").write_text(json.dumps({"meta": state["meta"]}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    for source in sorted(DATA_DIR.iterdir()):
        if source.is_file():
            shutil.copy2(source, ENV_DIR / source.name)


if __name__ == "__main__":
    main()
