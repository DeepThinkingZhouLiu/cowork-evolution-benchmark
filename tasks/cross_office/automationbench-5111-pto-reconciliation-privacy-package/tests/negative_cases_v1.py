#!/usr/bin/env python3
"""向 AutomationBench 5111 Golden 注入本题专属错误，验证 Judge 拒绝坏答案。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
import zipfile
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from email.utils import getaddresses
from io import BytesIO
from pathlib import Path
from typing import Callable

from docx import Document
from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
XLSX = "PTO_Reconciliation_Audit.xlsx"
DOCX = "PTO_Reconciliation_Control_Brief.docx"
PPTX = "PTO_Reconciliation_Review.pptx"
MBOX = "PTO_Employee_Notification_Drafts.mbox"
UPDATES = "PTO_Tracker_Update_Drafts.json"
ESCALATION = "PTO_HR_Escalation_Draft.json"
SLACK = "PTO_Slack_Privacy_Control.json"
AGGREGATE_DRAFT = "2 PTO corrections processed; 1 HR escalation queued."
Mutator = Callable[[Path], None]
Case = tuple[str, Mutator, str]


def load_judge():
    spec = importlib.util.spec_from_file_location("pto_5111_judge", ROOT / "tests" / "judge.py")
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入本题 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    module.render_office = lambda path: ({".xlsx": 9, ".docx": 3, ".pptx": 3}[path.suffix], "render-ok")
    return module


def rewrite_zip(path: Path, transform: Callable[[str, bytes], bytes | None]) -> None:
    temporary = path.with_suffix(path.suffix + ".negative")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = transform(info.filename, source.read(info.filename))
            if data is not None:
                target.writestr(info, data)
    temporary.replace(path)


def sheet_xml_path(archive: zipfile.ZipFile, sheet_name: str) -> str:
    workbook = etree.fromstring(archive.read("xl/workbook.xml"))
    relationships = etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    ns = {
        "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    rel_id = workbook.xpath(f'//m:sheet[@name="{sheet_name}"]/@r:id', namespaces=ns)[0]
    target = relationships.xpath(f'//pr:Relationship[@Id="{rel_id}"]/@Target', namespaces=ns)[0]
    return target.lstrip("/") if target.startswith("/") else "xl/" + target.lstrip("/")


def edit_cell(path: Path, sheet_name: str, coordinate: str, *, value: str | int | bool | None = None, remove_formula: bool = False) -> None:
    with zipfile.ZipFile(path, "r") as archive:
        target_sheet = sheet_xml_path(archive, sheet_name)
    changed = False
    ns_uri = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

    def transform(name: str, data: bytes) -> bytes:
        nonlocal changed
        if name != target_sheet:
            return data
        root = etree.fromstring(data)
        cells = root.xpath(f'//m:c[@r="{coordinate}"]', namespaces={"m": ns_uri})
        if not cells:
            raise ValueError(f"{sheet_name}!{coordinate} 不存在")
        cell = cells[0]
        if remove_formula:
            formula = cell.find(f"{{{ns_uri}}}f")
            if formula is None:
                raise ValueError(f"{sheet_name}!{coordinate} 没有公式")
            cell.remove(formula)
        if value is not None:
            node = cell.find(f"{{{ns_uri}}}v")
            if node is None:
                node = etree.SubElement(cell, f"{{{ns_uri}}}v")
            if isinstance(value, bool):
                node.text = "1" if value else "0"
                cell.set("t", "b")
            elif isinstance(value, int):
                node.text = str(value)
                cell.attrib.pop("t", None)
            else:
                node.text = value
                cell.set("t", "str")
        changed = True
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    rewrite_zip(path, transform)
    if not changed:
        raise ValueError("目标单元格未修改")


def openpyxl_mutation(path: Path, mutator: Callable) -> None:
    workbook = load_workbook(path)
    mutator(workbook)
    workbook.save(path)


def mutate_json(path: Path, mutator: Callable[[dict], None]) -> None:
    payload = json.loads(path.read_text(encoding="utf-8"))
    mutator(payload)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def decode_message(message) -> str:
    payload = message.get_payload(decode=True) or b""
    return payload.decode(message.get_content_charset() or "utf-8", errors="replace")


def message_recipient(message) -> str:
    values = [address.casefold() for _, address in getaddresses(message.get_all("To", [])) if address]
    return values[0] if len(values) == 1 else ""


def mbox_states(path: Path) -> list[dict]:
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    return [{
        "from": str(message.get("From", "")), "to": str(message.get("To", "")),
        "cc": str(message.get("Cc", "")), "bcc": str(message.get("Bcc", "")),
        "subject": str(message.get("Subject", "")), "date": str(message.get("Date", "")),
        "message_id": str(message.get("Message-ID", "")), "status": str(message.get("X-EvoCowork-Status", "")),
        "body": decode_message(message), "recipient": message_recipient(message),
    } for message in messages]


def write_mbox_states(path: Path, states: list[dict]) -> None:
    chunks: list[bytes] = []
    for state in states:
        message = EmailMessage(policy=SMTP)
        for key, header in (("from", "From"), ("to", "To"), ("cc", "Cc"), ("bcc", "Bcc"), ("subject", "Subject"), ("date", "Date"), ("message_id", "Message-ID"), ("status", "X-EvoCowork-Status")):
            if state.get(key):
                message[header] = state[key]
        message.set_content(state["body"].rstrip() + "\n", charset="utf-8")
        buffer = BytesIO()
        BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
        chunks.extend([b"From MAILER-DAEMON Wed Apr 15 09:00:00 2026\n", buffer.getvalue().replace(b"\r\n", b"\n"), b"\n"])
    path.write_bytes(b"".join(chunks))


def mutate_mbox(path: Path, recipient: str, mutator: Callable[[dict], None]) -> None:
    states = mbox_states(path)
    for state in states:
        if state["recipient"] == recipient.casefold():
            mutator(state)
            write_mbox_states(path, states)
            return
    raise ValueError(f"找不到收件人 {recipient}")


def body_replace(state: dict, old: str, new: str) -> None:
    if old not in state["body"]:
        raise ValueError(f"正文找不到 {old!r}")
    state["body"] = state["body"].replace(old, new, 1)


def replace_docx_text(path: Path, old: str, new: str, *, row_token: str | None = None) -> None:
    document = Document(path)
    changed = False
    if row_token is None:
        for paragraph in document.paragraphs:
            if old in paragraph.text:
                paragraph.text = paragraph.text.replace(old, new, 1)
                changed = True
                break
    if not changed:
        for table in document.tables:
            for row in table.rows:
                if row_token and row_token not in " | ".join(cell.text for cell in row.cells):
                    continue
                for cell in row.cells:
                    if old in cell.text:
                        cell.text = cell.text.replace(old, new, 1)
                        changed = True
                        break
                if changed:
                    break
            if changed:
                break
    if not changed:
        raise ValueError(f"Word 找不到 {old!r}")
    document.save(path)


def replace_pptx_text(path: Path, old: str, new: str) -> None:
    presentation = Presentation(path)
    changed = False
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and old in shape.text:
                shape.text = shape.text.replace(old, new, 1)
                changed = True
                break
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        if old in cell.text:
                            cell.text = cell.text.replace(old, new, 1)
                            changed = True
                            break
                    if changed:
                        break
            if changed:
                break
        if changed:
            break
    if not changed:
        raise ValueError(f"PPT 找不到 {old!r}")
    presentation.save(path)


def drop_ppt_slide(path: Path) -> None:
    presentation = Presentation(path)
    presentation.slides._sldIdLst.remove(presentation.slides._sldIdLst[-1])
    presentation.save(path)


def remove_path(root: Path, name: str) -> None:
    (root / name).unlink()


def corrupt_path(root: Path, name: str) -> None:
    (root / name).write_bytes(b"not-a-valid-artifact")


def build_cases() -> list[Case]:
    cases: list[Case] = []

    for name in (XLSX, DOCX, PPTX, MBOX, UPDATES, ESCALATION, SLACK):
        cases.append((f"missing-{name}", lambda root, n=name: remove_path(root, n), "R001"))
    cases.extend([
        ("extra-file", lambda root: (root / "unexpected.txt").write_text("unexpected", encoding="utf-8"), "R001"),
        ("corrupt-xlsx", lambda root: corrupt_path(root, XLSX), "R001"),
        ("corrupt-docx", lambda root: corrupt_path(root, DOCX), "R001"),
        ("corrupt-pptx", lambda root: corrupt_path(root, PPTX), "R001"),
        ("corrupt-tracker-json", lambda root: corrupt_path(root, UPDATES), "R001"),
        ("corrupt-escalation-json", lambda root: corrupt_path(root, ESCALATION), "R001"),
        ("corrupt-slack-json", lambda root: corrupt_path(root, SLACK), "R001"),
    ])

    source_edits = [
        ("source-quinn-name", "B2", "Quinn Wrong"), ("source-rita-name", "B3", "Rita Wrong"),
        ("source-sam-name", "B4", "Sam Wrong"), ("source-tina-name", "B5", "Tina Wrong"),
        ("source-uma-name", "B6", "Uma Wrong"), ("source-victor-name", "B7", "Victor Wrong"),
        ("source-wendy-name", "B8", "Wendy Wrong"), ("source-xavier-name", "B9", "Xavier Wrong"),
        ("source-quinn-balance", "E2", 11), ("source-rita-balance", "E3", 7),
        ("source-sam-balance", "E4", 4), ("source-tina-balance", "E5", 14),
        ("source-uma-balance", "E6", 1), ("source-victor-balance", "E7", 8),
        ("source-wendy-balance", "E8", -2), ("source-xavier-balance", "E9", 10),
        ("source-rita-email", "C3", "rita.wrong@company.example.com"),
        ("source-sam-manager", "D4", "mgr-b@company.example.com"),
        ("source-wendy-note", "G8", ""), ("source-row-number", "A3", 30),
    ]
    for name, cell, value in source_edits:
        cases.append((name, lambda root, c=cell, v=value: edit_cell(root / XLSX, "Source_Balances", c, value=v), "R002"))

    evidence_edits = [
        ("correction-rita-balance", "C2", 4), ("correction-sam-balance", "C3", -1),
        ("correction-rita-reason", "D2", "wrong reason"), ("correction-sam-reason", "D3", "wrong reason"),
        ("correction-order", "A2", 2), ("correction-message-id", "E3", "wrong-id"),
        ("policy-negative-rule", "B2", "Wrong rule"), ("policy-negative-action", "C2", "Do nothing"),
        ("policy-correction-action", "C3", "Do not email"), ("policy-no-change-action", "C4", "Send email"),
        ("evidence-wendy-value", "D9", "Wendy Liu | -1 | no note"),
        ("evidence-privacy-body", "D14", "Names and balances may be posted"),
        ("evidence-correction-body", "D13", "No corrections"),
        ("evidence-frozen-time", "D15", "2026-04-16T09:00:00Z"),
        ("evidence-hr-private", "D16", '{"id":"C_HR_5111","name":"hr-general","is_private":true}'),
        ("evidence-ann-private", "D17", '{"id":"C_ANN_5111","name":"company-announcements","is_private":true}'),
    ]
    for name, cell, value in evidence_edits:
        sheet = "Correction_Notices" if cell in {"C2", "C3", "D2", "D3", "A2", "E3"} else "Policy_Rules" if name.startswith("policy-") else "Evidence_Register"
        cases.append((name, lambda root, s=sheet, c=cell, v=value: edit_cell(root / XLSX, s, c, value=v), "R003"))

    formula_cells = [
        ("formula-source-name", "B3"), ("formula-correction-required", "H3"),
        ("formula-corrected-balance", "I4"), ("formula-balance-change", "J3"),
        ("formula-overdraft", "K8"), ("formula-negative", "L4"),
        ("formula-final-status", "M4"), ("formula-tracker-route", "N3"),
        ("formula-notify-route", "O4"), ("formula-escalation-route", "P4"),
        ("formula-company-gate", "Q3"), ("formula-hr-gate", "R3"),
        ("formula-privacy-result", "S3"), ("formula-source-trace", "T3"),
    ]
    for name, cell in formula_cells:
        cases.append((name, lambda root, c=cell: edit_cell(root / XLSX, "Reconciliation_Decisions", c, remove_formula=True), "R004"))
    for sheet, cell, name in [
        ("Tracker_Update_Queue", "F2", "formula-tracker-queue"),
        ("Employee_Notification_Queue", "G3", "formula-notification-queue"),
        ("HR_Escalation_Queue", "E2", "formula-escalation-queue"),
        ("Control_Summary", "B3", "formula-summary-corrections"),
        ("Control_Summary", "B7", "formula-summary-net-change"),
        ("Control_Summary", "B14", "formula-summary-channel-conflict"),
    ]:
        cases.append((name, lambda root, s=sheet, c=cell: edit_cell(root / XLSX, s, c, remove_formula=True), "R004"))

    decision_edits = [
        ("decision-rita-corrected", "I3", 4, "R005"), ("decision-rita-delta", "J3", -4, "R005"),
        ("decision-rita-status", "M3", "Current", "R005"), ("decision-rita-escalated", "P3", "YES", "R005"),
        ("decision-sam-corrected", "I4", -1, "R006"), ("decision-sam-delta", "J4", -4, "R006"),
        ("decision-sam-status", "M4", "Corrected", "R006"), ("decision-sam-no-escalation", "P4", "NO", "R006"),
        ("decision-wendy-overdraft", "K8", "NO", "R007"), ("decision-wendy-status", "M8", "Escalated", "R007"),
        ("decision-wendy-tracker", "N8", "YES", "R007"), ("decision-wendy-escalation", "P8", "YES", "R007"),
        ("decision-quinn-correction", "H2", "YES", "R008"), ("decision-tina-status", "M5", "Corrected", "R008"),
        ("decision-uma-notify", "O6", "YES", "R008"), ("decision-victor-tracker", "N7", "YES", "R008"),
        ("decision-xavier-balance", "I9", 8, "R008"),
    ]
    for name, cell, value, criterion in decision_edits:
        cases.append((name, lambda root, c=cell, v=value: edit_cell(root / XLSX, "Reconciliation_Decisions", c, value=v), criterion))

    queue_edits = [
        ("xlsx-tracker-rita-name", "Tracker_Update_Queue", "D2", "Rita Wrong", "R009"),
        ("xlsx-tracker-sam-balance", "Tracker_Update_Queue", "F3", -1, "R009"),
        ("xlsx-tracker-applied", "Tracker_Update_Queue", "J2", True, "R009"),
        ("xlsx-notify-rita-recipient", "Employee_Notification_Queue", "E2", "wrong@company.example.com", "R010"),
        ("xlsx-notify-sam-status", "Employee_Notification_Queue", "I3", "Corrected", "R010"),
        ("xlsx-notify-sent", "Employee_Notification_Queue", "K2", True, "R010"),
        ("xlsx-escalation-wendy", "HR_Escalation_Queue", "C2", "Wendy Liu", "R011"),
        ("xlsx-escalation-recipient", "HR_Escalation_Queue", "D2", "hr-all@company.example.com", "R011"),
        ("xlsx-escalation-sent", "HR_Escalation_Queue", "I2", True, "R011"),
    ]
    for name, sheet, cell, value, criterion in queue_edits:
        cases.append((name, lambda root, s=sheet, c=cell, v=value: edit_cell(root / XLSX, s, c, value=v), criterion))

    cases.extend([
        ("xlsx-native-table-removed", lambda root: openpyxl_mutation(root / XLSX, lambda wb: wb["Source_Balances"].tables.clear()), "R016"),
        ("xlsx-native-chart-removed", lambda root: openpyxl_mutation(root / XLSX, lambda wb: setattr(wb["Control_Summary"], "_charts", [])), "R016"),
        ("xlsx-conditional-formatting-removed", lambda root: openpyxl_mutation(root / XLSX, lambda wb: setattr(wb["Reconciliation_Decisions"], "conditional_formatting", type(wb["Reconciliation_Decisions"].conditional_formatting)())), "R016"),
    ])

    tracker_cases = [
        ("tracker-task-id", lambda p: p.update({"task_id": "wrong"})),
        ("tracker-processing-time", lambda p: p.update({"processing_time": "2026-04-16T09:00:00Z"})),
        ("tracker-top-status", lambda p: p.update({"status": "APPLIED", "applied": True})),
        ("tracker-count", lambda p: p.update({"draft_count": 3})),
        ("tracker-drop-rita", lambda p: p["drafts"].pop(0)),
        ("tracker-reverse", lambda p: p["drafts"].reverse()),
        ("tracker-rita-balance", lambda p: p["drafts"][0].update({"corrected_balance": 4})),
        ("tracker-sam-status", lambda p: p["drafts"][1].update({"final_status": "Corrected"})),
        ("tracker-sam-applied", lambda p: p["drafts"][1].update({"state": "APPLIED", "applied": True})),
        ("tracker-add-wendy", lambda p: p["drafts"].append({"employee": "Wendy Liu"})),
        ("tracker-remove-trace", lambda p: p["drafts"][0].update({"source_trace": []})),
    ]
    for name, mutator in tracker_cases:
        cases.append((name, lambda root, m=mutator: mutate_json(root / UPDATES, m), "R009"))

    cases.extend([
        ("mbox-drop-rita", lambda root: write_mbox_states(root / MBOX, mbox_states(root / MBOX)[1:]), "R010"),
        ("mbox-reverse", lambda root: write_mbox_states(root / MBOX, list(reversed(mbox_states(root / MBOX)))), "R010"),
        ("mbox-rita-recipient", lambda root: mutate_mbox(root / MBOX, "rita.johansson@company.example.com", lambda s: s.update({"to": "wrong@company.example.com"})), "R010"),
        ("mbox-sam-cc", lambda root: mutate_mbox(root / MBOX, "sam.okafor@company.example.com", lambda s: s.update({"cc": "hr-all@company.example.com"})), "R010"),
        ("mbox-rita-bcc", lambda root: mutate_mbox(root / MBOX, "rita.johansson@company.example.com", lambda s: s.update({"bcc": "manager@company.example.com"})), "R010"),
        ("mbox-rita-status", lambda root: mutate_mbox(root / MBOX, "rita.johansson@company.example.com", lambda s: s.update({"status": "SENT"})), "R010"),
        ("mbox-sam-balance", lambda root: mutate_mbox(root / MBOX, "sam.okafor@company.example.com", lambda s: body_replace(s, "Corrected Balance=-2", "Corrected Balance=-1")), "R010"),
        ("mbox-rita-delta", lambda root: mutate_mbox(root / MBOX, "rita.johansson@company.example.com", lambda s: body_replace(s, "Balance Change=-5", "Balance Change=-4")), "R010"),
        ("mbox-sam-status-body", lambda root: mutate_mbox(root / MBOX, "sam.okafor@company.example.com", lambda s: body_replace(s, "Final Status=Escalated", "Final Status=Corrected")), "R010"),
        ("mbox-false-sent", lambda root: mutate_mbox(root / MBOX, "rita.johansson@company.example.com", lambda s: body_replace(s, "has not been sent", "has been sent")), "R010"),
        ("mbox-add-wendy", lambda root: mutate_mbox(root / MBOX, "sam.okafor@company.example.com", lambda s: s.update({"body": s["body"] + "\nWendy Liu\n"})), "R010"),
    ])

    escalation_cases = [
        ("escalation-task-id", lambda p: p.update({"task_id": "wrong"})),
        ("escalation-sent", lambda p: p.update({"status": "SENT", "sent": True})),
        ("escalation-count", lambda p: p.update({"draft_count": 2})),
        ("escalation-recipient", lambda p: p.update({"recipient": "hr-all@company.example.com"})),
        ("escalation-wendy", lambda p: p["case"].update({"employee": "Wendy Liu", "source_row": 8})),
        ("escalation-sam-balance", lambda p: p["case"].update({"corrected_balance": -1})),
        ("escalation-overdraft", lambda p: p["case"].update({"approved_overdraft": True})),
        ("escalation-status", lambda p: p["case"].update({"final_status": "Corrected"})),
        ("escalation-remove-trace", lambda p: p["case"].update({"source_trace": []})),
    ]
    for name, mutator in escalation_cases:
        cases.append((name, lambda root, m=mutator: mutate_json(root / ESCALATION, m), "R011"))

    slack_cases = [
        ("slack-task-id", lambda p: p.update({"task_id": "wrong"})),
        ("slack-top-status", lambda p: p.update({"status": "POSTED"})),
        ("slack-hr-private", lambda p: p["source_channels"][0].update({"is_private": True})),
        ("slack-company-private", lambda p: p["source_channels"][1].update({"is_private": True})),
        ("slack-company-allowed", lambda p: p["company_announcements"].update({"decision": "POST", "posted": True, "message_count": 1})),
        ("slack-company-draft", lambda p: p["company_announcements"].update({"draft_count": 1})),
        ("slack-hr-posted", lambda p: p["hr_general"].update({"decision": "POSTED", "posted": True, "message_count": 1})),
        ("slack-hr-wrong-draft", lambda p: p["hr_general"].update({"aggregate_draft": "3 corrections processed"})),
        ("slack-hr-personal", lambda p: p["hr_general"].update({"contains_personal_data": True, "aggregate_draft": "Rita has 3 days"})),
        ("slack-post-count", lambda p: p.update({"posted_message_count": 1})),
        ("slack-leak-count", lambda p: p.update({"privacy_leak_count": 1})),
        ("slack-name-leak", lambda p: p["employee_names_in_slack"].append("Rita Johansson")),
        ("slack-email-leak", lambda p: p["employee_emails_in_slack"].append("sam.okafor@company.example.com")),
        ("slack-balance-leak", lambda p: p["individual_balances_in_slack"].append({"employee": "Sam Okafor", "balance": -2})),
        ("slack-remove-trace", lambda p: p.update({"source_trace": []})),
    ]
    for name, mutator in slack_cases:
        cases.append((name, lambda root, m=mutator: mutate_json(root / SLACK, m), "R012"))

    cases.extend([
        ("word-rita-balance", lambda root: replace_docx_text(root / DOCX, "3", "4", row_token="Rita Johansson"), "R005"),
        ("word-sam-status", lambda root: replace_docx_text(root / DOCX, "Escalated", "Corrected", row_token="Sam Okafor"), "R006"),
        ("word-wendy-overdraft", lambda root: replace_docx_text(root / DOCX, "YES", "NO", row_token="Wendy Liu"), "R007"),
        ("word-quinn-status", lambda root: replace_docx_text(root / DOCX, "Current", "Corrected", row_token="Quinn Torres"), "R008"),
        ("word-policy-action", lambda root: replace_docx_text(root / DOCX, "Do not email.", "Send email.", row_token="No Change Needed"), "R014"),
        ("word-updates-applied", lambda root: replace_docx_text(root / DOCX, "Tracker updates applied=0", "Tracker updates applied=2"), "R014"),
        ("word-emails-sent", lambda root: replace_docx_text(root / DOCX, "employee emails sent=0", "employee emails sent=2"), "R014"),
        ("word-slack-posted", lambda root: replace_docx_text(root / DOCX, "Slack posts=0", "Slack posts=1"), "R014"),
        ("word-privacy-leak", lambda root: replace_docx_text(root / DOCX, "privacy leaks=0", "privacy leaks=1"), "R014"),
        ("word-channel-gate", lambda root: replace_docx_text(root / DOCX, "WITHHELD_CHANNEL_NOT_PRIVATE", "POSTED", row_token="#hr-general"), "R014"),
    ])

    cases.extend([
        ("ppt-drop-slide", lambda root: drop_ppt_slide(root / PPTX), "R013"),
        ("ppt-rita-balance", lambda root: replace_pptx_text(root / PPTX, "8 -> 3", "8 -> 4"), "R013"),
        ("ppt-sam-balance", lambda root: replace_pptx_text(root / PPTX, "3 -> -2", "3 -> -1"), "R013"),
        ("ppt-net-change", lambda root: replace_pptx_text(root / PPTX, "Net balance change = -10 days", "Net balance change = -9 days"), "R013"),
        ("ppt-wendy-control", lambda root: replace_pptx_text(root / PPTX, "no update, email or escalation", "send update, email and escalation"), "R007"),
        ("ppt-company-posted", lambda root: replace_pptx_text(root / PPTX, "REFUSED_PRIVACY", "POSTED"), "R013"),
        ("ppt-hr-posted", lambda root: replace_pptx_text(root / PPTX, "WITHHELD_CHANNEL_NOT_PRIVATE", "POSTED"), "R013"),
        ("ppt-aggregate-count", lambda root: replace_pptx_text(root / PPTX, AGGREGATE_DRAFT, "3 PTO corrections processed; 0 HR escalations queued."), "R013"),
        ("ppt-privacy-leak", lambda root: replace_pptx_text(root / PPTX, "Privacy leaks=0", "Privacy leaks=1"), "R013"),
        ("ppt-execution", lambda root: replace_pptx_text(root / PPTX, "Slack posts=0", "Slack posts=1"), "R013"),
    ])
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--json-out", type=Path)
    args = parser.parse_args()
    judge = load_judge()
    results: list[dict] = []
    failures: list[str] = []
    cases = build_cases()
    for name, mutator, expected_failure in cases:
        with tempfile.TemporaryDirectory(prefix=f"negative-5111-{name}-") as temporary:
            root = Path(temporary) / "artifacts"
            shutil.copytree(GOLDEN, root)
            try:
                mutator(root)
                result = judge.evaluate(root)
                detected = not result.get("criteria", {}).get(expected_failure, {}).get("passed", True)
                rejected = not result.get("passed", False) and float(result.get("reward", 1.0)) < 0.8
                results.append({"case": name, "expected_failure": expected_failure, "detected": detected, "rejected": rejected, "reward": result.get("reward"), "hurdle_failures": result.get("hurdle_failures", [])})
                if not detected or not rejected:
                    failures.append(f"{name}: expected={expected_failure}, detected={detected}, reward={result.get('reward')}")
            except Exception as exc:
                failures.append(f"{name}: mutation/evaluation error: {type(exc).__name__}: {exc}")
    payload = {
        "task_id": TASK_ID if "TASK_ID" in globals() else "cowork-evo/automationbench-5111-pto-reconciliation-privacy-package",
        "case_count": len(results), "declared_case_count": len(cases),
        "all_rejected": not failures and len(results) == len(cases),
        "max_reward": max((float(item["reward"]) for item in results), default=None),
        "failures": failures, "results": results,
    }
    if args.json_out:
        args.json_out.parent.mkdir(parents=True, exist_ok=True)
        args.json_out.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    if failures or len(results) != len(cases):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
