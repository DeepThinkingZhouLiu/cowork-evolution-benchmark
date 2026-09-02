#!/usr/bin/env python3
"""只把 AutomationBench 5112 的公开初始状态静态化为本题输入。

本脚本只接受 ``hr.promotion_announcement``，逐项断言八条候选记录、
两封完整邮件、冻结时间、Slack 频道和八条上游断言。它不是通用制题器，
不会创造员工、职级、部门、备注、经理身份或公告规则。
"""

from __future__ import annotations

import hashlib
import importlib.util
import json
import shutil
import sys
import types
import zipfile
from datetime import datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
UPSTREAM = Path(__file__).with_name("upstream-hr-tasks.py")
SNAPSHOT = Path(__file__).with_name("upstream-task-5112.json")
LICENSE = Path(__file__).with_name("LICENSE-AutomationBench-MIT")
DATA_DIR = ROOT / "data" / "input_files"
ENV_DIR = ROOT / "environment" / "input_files"
FIXED_ZIP_TIME = (2026, 4, 15, 9, 0, 0)
SOURCE_SHA256 = "c173a36a03f43096bd122577584aa32f9df379d8760aa46c43be04c85d101779"
LICENSE_SHA256 = "c6f9d87aff6d627d5cd78cb09364dafed7d6f1c4bb7a91183d57377157e96ba4"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def freeze_upstream_assets() -> None:
    source_task = ROOT.parent / "automationbench-5111-pto-reconciliation-privacy-package" / "adaptation"
    if not UPSTREAM.is_file():
        source = source_task / "upstream-hr-tasks.py"
        if not source.is_file() or sha256(source) != SOURCE_SHA256:
            raise FileNotFoundError("找不到哈希匹配的 AutomationBench HR 冻结源文件")
        shutil.copy2(source, UPSTREAM)
    if sha256(UPSTREAM) != SOURCE_SHA256:
        raise ValueError("AutomationBench HR 冻结源文件哈希错误")
    if not LICENSE.is_file():
        source = source_task / "LICENSE-AutomationBench-MIT"
        if not source.is_file() or sha256(source) != LICENSE_SHA256:
            raise FileNotFoundError("找不到哈希匹配的 AutomationBench LICENSE")
        shutil.copy2(source, LICENSE)
    if sha256(LICENSE) != LICENSE_SHA256:
        raise ValueError("AutomationBench LICENSE 哈希错误")


def load_task() -> dict:
    datasets = types.ModuleType("datasets")
    datasets.Dataset = object
    sys.modules["datasets"] = datasets
    for name in ("automationbench", "automationbench.domains", "automationbench.domains.hr"):
        sys.modules[name] = types.ModuleType(name)
    noise = types.ModuleType("automationbench.domains.hr._noise")
    noise.apply_noise = lambda value: value
    sys.modules[noise.__name__] = noise
    spec = importlib.util.spec_from_file_location("automationbench_hr_tasks_frozen_5112", UPSTREAM)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入冻结的 AutomationBench HR 源文件")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    task = module.get_task_5112()
    if task.get("example_id") != 5112 or task.get("task") != "hr.promotion_announcement":
        raise ValueError("冻结函数不再是 AutomationBench 5112")
    return task


def assert_shape(task: dict) -> None:
    prompt = task["prompt"][1]["content"]
    assert all(token in prompt for token in ("Q1 promotion announcements", "including the lateral moves", "Same format as last quarter", "announcement guidelines"))
    state = task["info"]["initial_state"]
    assert state["meta"] == {"schema_version": "0.1.0", "current_time": "2026-04-15T09:00:00Z"}
    messages = state["gmail"]["messages"]
    assert [message["id"] for message in messages] == ["msg_last_promo_5112", "msg_announce_policy_5112"]
    assert all(token in messages[0]["body_plain"] for token in ("Jennifer Park", "Staff Engineer", "Michael Adebayo", "Senior Account Manager"))
    assert all(token in messages[1]["body_plain"] for token in ("Only vertical promotions", "Lateral moves", "NOT included", "all-company@company.example.com", "hold", "pending"))
    sheet = state["google_sheets"]["spreadsheets"][0]["worksheets"][0]
    rows = sheet["rows"]
    assert [row["row_id"] for row in rows] == list(range(2, 10))
    assert [row["cells"]["Employee"] for row in rows] == ["Ava Chen", "Brandon Osei", "Clara Dubois", "Derek Holm", "Elena Vasquez", "Faisal Rahman", "Greta Nilsson", "Hugo Fernandez"]
    assert [row["cells"]["Type"] for row in rows] == ["Vertical", "Vertical", "Lateral", "Vertical", "Vertical", "Lateral", "Vertical", "Vertical"]
    assert rows[2]["cells"]["Notes"] == "Moving from Growth to Platform team"
    assert rows[3]["cells"]["Notes"] == "Pending - awaiting final sign-off from CTO"
    assert rows[5]["cells"]["Notes"] == "Transfer to EMEA team"
    assert state["slack"] == {"channels": [{"id": "C_GEN_5112", "name": "general", "is_private": False}], "messages": []}
    assert len(task["info"]["assertions"]) == 8


def normalize_xlsx(path: Path) -> None:
    temporary = path.with_suffix(".normalized.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    temporary.replace(path)


def style_sheet(sheet, widths: dict[str, float]) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = sheet.dimensions


def write_roster(state: dict, target: Path) -> None:
    source = state["google_sheets"]["spreadsheets"][0]
    worksheet = source["worksheets"][0]
    workbook = Workbook()
    workbook.creator = "AutomationBench 5112 public fixture"
    workbook.title = source["title"]
    workbook.subject = "Static Office representation of the Q1 promotion roster"
    workbook.description = "No employee, title, department, type or note was added beyond AutomationBench 5112."
    sheet = workbook.active
    sheet.title = worksheet["title"]
    sheet.append(["Source Row", "Employee", "Department", "Previous Title", "New Title", "Type", "Notes"])
    for item in worksheet["rows"]:
        cells = item["cells"]
        sheet.append([item["row_id"], cells["Employee"], cells["Department"], cells["Previous Title"], cells["New Title"], cells["Type"], cells["Notes"]])
    table = Table(displayName="PromotionRosterSourceInput", ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    style_sheet(sheet, {"A": 12, "B": 24, "C": 22, "D": 30, "E": 32, "F": 15, "G": 52})
    workbook.save(target)
    normalize_xlsx(target)


def message_bytes(record: dict) -> bytes:
    message = EmailMessage(policy=SMTP)
    message["From"] = record["from_"]
    message["To"] = ", ".join(record["to"])
    message["Subject"] = record["subject"]
    parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
    message["Date"] = parsed.strftime("%a, %d %b %Y %H:%M:%S +0000")
    message["Message-ID"] = f"<{record['id']}@automationbench.invalid>"
    message["X-AutomationBench-Message-ID"] = record["id"]
    message["X-AutomationBench-Thread-ID"] = record["thread_id"]
    message["X-AutomationBench-Labels"] = ",".join(record.get("label_ids", []))
    message.set_content(record["body_plain"], charset="utf-8")
    buffer = BytesIO()
    BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
    return buffer.getvalue()


def write_mbox(messages: list[dict], target: Path) -> None:
    chunks: list[bytes] = []
    for record in messages:
        parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
        chunks.append(f"From {record['from_']} {parsed.strftime('%a %b %d %H:%M:%S %Y')}\n".encode("ascii"))
        chunks.append(message_bytes(record).replace(b"\r\n", b"\n"))
        chunks.append(b"\n")
    target.write_bytes(b"".join(chunks))


def main() -> None:
    freeze_upstream_assets()
    task = load_task()
    assert_shape(task)
    state = task["info"]["initial_state"]
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    ENV_DIR.mkdir(parents=True, exist_ok=True)
    SNAPSHOT.write_text(json.dumps(task, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_roster(state, DATA_DIR / "Q1_Promotion_Roster.xlsx")
    write_mbox(state["gmail"]["messages"], DATA_DIR / "Promotion_Announcement_Context.mbox")
    (DATA_DIR / "Slack_Channels.json").write_text(json.dumps(state["slack"], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (DATA_DIR / "Workflow_Context.json").write_text(json.dumps({"meta": state["meta"]}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    for source in sorted(DATA_DIR.iterdir()):
        if source.is_file():
            shutil.copy2(source, ENV_DIR / source.name)


if __name__ == "__main__":
    main()
