#!/usr/bin/env python3
"""AutomationBench 5112 晋升公告治理包的逐题参考解法。"""

from __future__ import annotations

import argparse
import hashlib
import json
import mailbox
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt as PptPt


INPUT_HASHES = {
    "Q1_Promotion_Roster.xlsx": "3c0b8c890197cfabc12a0b0f6fac0e9620898fc6c66d6bc10be9fc9c8062a6cd",
    "Promotion_Announcement_Context.mbox": "f7389ea8ddce77ea6c3bfbfb8b0951fdcf0ffbfa967854de622fdb8251f3b610",
    "Slack_Channels.json": "839f793bd3a10a9487d92c86b38d923c90e9ff4015f57c6e0dee3de589287ec5",
    "Workflow_Context.json": "76c4a0657bc27f2fd830b10613ec7ceae3adcdc5798aeee0f822d27bf69bbc99",
}
XLSX_NAME = "Promotion_Eligibility_Audit.xlsx"
DOCX_NAME = "Promotion_Announcement_Control_Brief.docx"
PPTX_NAME = "Promotion_Announcement_Review.pptx"
MBOX_NAME = "Company_Promotion_Announcement_Draft.mbox"
EXCLUSION_NAME = "Promotion_Exclusion_Register.json"
SLACK_NAME = "Promotion_Slack_Control.json"
EXECUTION_NAME = "Promotion_Execution_Control.json"
FIXED_ISO_TIME = "2026-04-15T09:00:00Z"
FIXED_ZIP_TIME = (2026, 4, 15, 9, 0, 0)
SUBJECT = "Q1 2026 Promotions - Congratulations!"

NAVY = "17365D"; BLUE = "5B9BD5"; GREEN = "E2F0D9"; AMBER = "FFF2CC"; RED = "F4CCCC"; GRAY = "E7E6E6"; WHITE = "FFFFFF"
THIN = Side(style="thin", color="B7C9D6")
PPT_NAVY = PptRGBColor(23, 54, 93); PPT_BLUE = PptRGBColor(91, 155, 213); PPT_LIGHT_BLUE = PptRGBColor(221, 235, 247)
PPT_GREEN = PptRGBColor(226, 240, 217); PPT_AMBER = PptRGBColor(255, 242, 204); PPT_RED = PptRGBColor(244, 204, 204); PPT_GRAY = PptRGBColor(242, 242, 242); PPT_WHITE = PptRGBColor(255, 255, 255)


@dataclass(frozen=True)
class Promotion:
    source_row: int
    employee: str
    department: str
    previous_title: str
    new_title: str
    move_type: str
    notes: str


@dataclass(frozen=True)
class Decision:
    source: Promotion
    vertical: str
    lateral: str
    hold_pending: str
    company_eligible: str
    exclusion_reason: str
    company_email_route: str
    receiving_manager_route: str
    manager_identity: str
    manager_draft: str
    slack_route: str
    final_status: str


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def decode_message(message) -> str:
    if message.is_multipart():
        parts: list[str] = []
        for part in message.walk():
            if part.get_content_type() == "text/plain":
                payload = part.get_payload(decode=True) or b""
                parts.append(payload.decode(part.get_content_charset() or "utf-8", errors="replace"))
        return "\n".join(parts).strip()
    payload = message.get_payload(decode=True) or b""
    return payload.decode(message.get_content_charset() or "utf-8", errors="replace").strip()


def read_inputs(input_dir: Path) -> tuple[list[Promotion], dict[str, dict], dict, str]:
    for name, expected in INPUT_HASHES.items():
        source = input_dir / name
        if not source.is_file() or sha256(source) != expected:
            raise ValueError(f"输入 {name} 缺失或 SHA256 与冻结版本不一致")
    context = json.loads((input_dir / "Workflow_Context.json").read_text(encoding="utf-8"))
    if context != {"meta": {"schema_version": "0.1.0", "current_time": FIXED_ISO_TIME}}:
        raise ValueError("冻结运行时间已变化")
    channels = json.loads((input_dir / "Slack_Channels.json").read_text(encoding="utf-8"))
    if channels != {"channels": [{"id": "C_GEN_5112", "name": "general", "is_private": False}], "messages": []}:
        raise ValueError("Slack 频道快照已变化")

    workbook = load_workbook(input_dir / "Q1_Promotion_Roster.xlsx", read_only=True, data_only=True)
    if workbook.sheetnames != ["Approved Promotions"]:
        raise ValueError("晋升工作簿结构已变化")
    promotions: list[Promotion] = []
    for row in workbook["Approved Promotions"].iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        promotions.append(Promotion(int(row[0]), str(row[1]), str(row[2]), str(row[3]), str(row[4]), str(row[5]), str(row[6] or "")))
    workbook.close()
    expected_promotions = [
        (2, "Ava Chen", "Engineering", "Software Engineer", "Senior Software Engineer", "Vertical", ""),
        (3, "Brandon Osei", "Sales", "Account Executive", "Senior Account Executive", "Vertical", ""),
        (4, "Clara Dubois", "Product", "Product Manager", "Product Manager", "Lateral", "Moving from Growth to Platform team"),
        (5, "Derek Holm", "Engineering", "Senior Engineer", "Staff Engineer", "Vertical", "Pending - awaiting final sign-off from CTO"),
        (6, "Elena Vasquez", "Engineering", "Tech Lead", "Engineering Manager", "Vertical", ""),
        (7, "Faisal Rahman", "Marketing", "Marketing Analyst", "Marketing Analyst", "Lateral", "Transfer to EMEA team"),
        (8, "Greta Nilsson", "Support", "Support Specialist", "Senior Support Specialist", "Vertical", ""),
        (9, "Hugo Fernandez", "Engineering", "Junior Engineer", "Software Engineer", "Vertical", ""),
    ]
    if [(item.source_row, item.employee, item.department, item.previous_title, item.new_title, item.move_type, item.notes) for item in promotions] != expected_promotions:
        raise ValueError("八条冻结晋升记录已变化")

    box = mailbox.mbox(input_dir / "Promotion_Announcement_Context.mbox", create=False)
    source_messages = list(box)
    box.close()
    if len(source_messages) != 2:
        raise ValueError("公告上下文必须恰好两封邮件")
    messages: dict[str, dict] = {}
    for message in source_messages:
        message_id = str(message.get("X-AutomationBench-Message-ID", ""))
        messages[message_id] = {
            "message_id": message_id, "thread_id": str(message.get("X-AutomationBench-Thread-ID", "")),
            "from": str(message.get("From", "")), "to": str(message.get("To", "")),
            "subject": str(message.get("Subject", "")), "date": str(message.get("Date", "")),
            "body": decode_message(message),
        }
    if set(messages) != {"msg_last_promo_5112", "msg_announce_policy_5112"}:
        raise ValueError("权威公告邮件身份已变化")
    prior = messages["msg_last_promo_5112"]["body"]
    policy = messages["msg_announce_policy_5112"]["body"]
    if not all(token in prior for token in ("Jennifer Park", "Staff Engineer", "Michael Adebayo", "Senior Account Manager", "Congratulations to all!", "- HR Team")):
        raise ValueError("上季度公告格式证据不完整")
    if not all(token in policy for token in ("Only vertical promotions", "Lateral moves", "NOT included", "all-company@company.example.com", "hold", "pending", "receiving team's manager")):
        raise ValueError("VP People 公告政策证据不完整")
    return promotions, messages, channels, context["meta"]["current_time"]


def derive_decisions(promotions: list[Promotion]) -> list[Decision]:
    decisions: list[Decision] = []
    for source in promotions:
        vertical = "YES" if source.move_type == "Vertical" else "NO"
        lateral = "YES" if source.move_type == "Lateral" else "NO"
        notes = source.notes.casefold()
        hold_pending = "YES" if "hold" in notes or "pending" in notes else "NO"
        eligible = "YES" if vertical == "YES" and hold_pending == "NO" else "NO"
        if eligible == "YES":
            exclusion = "NONE_ELIGIBLE"
            manager_route = "NO"
            status = "ELIGIBLE_COMPANY_ANNOUNCEMENT"
        elif lateral == "YES":
            exclusion = "LATERAL_COMPANY_ANNOUNCEMENT_PROHIBITED"
            manager_route = "YES"
            status = "RECEIVING_MANAGER_COMMUNICATION_REQUIRED"
        else:
            exclusion = "PENDING_HOLD_NOT_CLEARED"
            manager_route = "NO"
            status = "PENDING_HOLD_NOT_CLEARED"
        decisions.append(Decision(source, vertical, lateral, hold_pending, eligible, exclusion, eligible, manager_route, "UNAVAILABLE_IN_INPUT" if manager_route == "YES" else "N/A", "NO", "NO", status))
    expected = {
        "Ava Chen": ("YES", "NO", "NO", "YES", "NONE_ELIGIBLE", "YES", "NO", "N/A", "NO", "NO", "ELIGIBLE_COMPANY_ANNOUNCEMENT"),
        "Brandon Osei": ("YES", "NO", "NO", "YES", "NONE_ELIGIBLE", "YES", "NO", "N/A", "NO", "NO", "ELIGIBLE_COMPANY_ANNOUNCEMENT"),
        "Clara Dubois": ("NO", "YES", "NO", "NO", "LATERAL_COMPANY_ANNOUNCEMENT_PROHIBITED", "NO", "YES", "UNAVAILABLE_IN_INPUT", "NO", "NO", "RECEIVING_MANAGER_COMMUNICATION_REQUIRED"),
        "Derek Holm": ("YES", "NO", "YES", "NO", "PENDING_HOLD_NOT_CLEARED", "NO", "NO", "N/A", "NO", "NO", "PENDING_HOLD_NOT_CLEARED"),
        "Elena Vasquez": ("YES", "NO", "NO", "YES", "NONE_ELIGIBLE", "YES", "NO", "N/A", "NO", "NO", "ELIGIBLE_COMPANY_ANNOUNCEMENT"),
        "Faisal Rahman": ("NO", "YES", "NO", "NO", "LATERAL_COMPANY_ANNOUNCEMENT_PROHIBITED", "NO", "YES", "UNAVAILABLE_IN_INPUT", "NO", "NO", "RECEIVING_MANAGER_COMMUNICATION_REQUIRED"),
        "Greta Nilsson": ("YES", "NO", "NO", "YES", "NONE_ELIGIBLE", "YES", "NO", "N/A", "NO", "NO", "ELIGIBLE_COMPANY_ANNOUNCEMENT"),
        "Hugo Fernandez": ("YES", "NO", "NO", "YES", "NONE_ELIGIBLE", "YES", "NO", "N/A", "NO", "NO", "ELIGIBLE_COMPANY_ANNOUNCEMENT"),
    }
    actual = {item.source.employee: (item.vertical, item.lateral, item.hold_pending, item.company_eligible, item.exclusion_reason, item.company_email_route, item.receiving_manager_route, item.manager_identity, item.manager_draft, item.slack_route, item.final_status) for item in decisions}
    if actual != expected:
        raise ValueError("八人公告资格真值不一致")
    return decisions


def add_table(sheet, name: str) -> None:
    table = Table(displayName=name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)


def style_sheet(sheet, widths: dict[str, float], footer: str, *, a3: bool = False) -> None:
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions; sheet.sheet_view.showGridLines = False
    for column, width in widths.items(): sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY); cell.font = Font(name="Aptos", bold=True, color=WHITE, size=8.3); cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = Border(bottom=THIN)
    sheet.row_dimensions[1].height = 31
    grid = Side(style="hair", color="D9E2F3")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=8.1); cell.alignment = Alignment(vertical="top", wrap_text=True); cell.border = Border(left=grid, right=grid, bottom=grid)
    sheet.page_setup.orientation = "landscape"; sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if a3 else sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 1; sheet.sheet_properties.pageSetUpPr.fitToPage = True; sheet.print_options.horizontalCentered = True
    sheet.oddFooter.center.text = footer; sheet.print_area = sheet.dimensions


def normalize_ooxml(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="promotion-5112-ooxml-") as temporary:
        target = Path(temporary) / path.name; axis_ids: dict[str, str] = {}
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as destination:
            for name in sorted(source.namelist()):
                data = source.read(name)
                if name == "docProps/core.xml":
                    try:
                        root = etree.fromstring(data); ns = {"dcterms": "http://purl.org/dc/terms/", "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"}
                        for tag in ("created", "modified"):
                            for node in root.xpath(f"//dcterms:{tag}", namespaces=ns): node.text = "2026-04-15T09:00:00Z"
                        for tag in ("lastModifiedBy", "revision"):
                            for node in root.xpath(f"//cp:{tag}", namespaces=ns): node.text = "EvoCowork" if tag == "lastModifiedBy" else "1"
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except etree.XMLSyntaxError: pass
                if name.startswith("xl/charts/chart") and name.endswith(".xml"):
                    try:
                        root = etree.fromstring(data); ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
                        for node in root.xpath("//c:axId", namespaces=ns):
                            old = node.get("val")
                            if old and old not in axis_ids: axis_ids[old] = str(51_120_001 + len(axis_ids))
                        for node in root.xpath("//c:axId | //c:crossAx", namespaces=ns):
                            old = node.get("val")
                            if old in axis_ids: node.set("val", axis_ids[old])
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except etree.XMLSyntaxError: pass
                info = zipfile.ZipInfo(name, FIXED_ZIP_TIME); info.compress_type = zipfile.ZIP_DEFLATED; info.external_attr = 0o600 << 16; destination.writestr(info, data)
        shutil.copy2(target, path)


def recalculate_xlsx(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="promotion-5112-recalc-") as temporary:
        root = Path(temporary); source_dir, output_dir, profile_dir = root / "source", root / "output", root / "profile"
        source_dir.mkdir(); output_dir.mkdir(); source = source_dir / path.name; shutil.copy2(path, source)
        process = subprocess.run(["libreoffice", "--headless", f"-env:UserInstallation=file://{profile_dir}", "--convert-to", "xlsx", "--outdir", str(output_dir), str(source)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300)
        converted = output_dir / path.name
        if process.returncode != 0 or not converted.is_file(): raise RuntimeError(f"LibreOffice 公式重算失败：{process.stdout[-1000:]}")
        shutil.copy2(converted, path)


def policy_rows() -> list[tuple[int, str, str]]:
    return [
        (1, "Company-wide scope", "Only vertical promotions (title level increase) are announced company-wide."),
        (2, "Lateral moves", "Lateral moves are NOT included; they are communicated separately by the receiving team's manager."),
        (3, "Company email route", "Announcements go to all-company@company.example.com."),
        (4, "Hold or pending", "Employees with a 'hold' or 'pending' note should not be announced until the note is cleared."),
    ]


def write_workbook(path: Path, promotions: list[Promotion], messages: dict[str, dict], channels: dict, frozen_time: str) -> None:
    workbook = Workbook(); workbook.remove(workbook.active)
    workbook.creator = "EvoCowork AutomationBench 5112 adaptation"; workbook.title = "Promotion Eligibility Audit"; workbook.subject = "Policy-governed company promotion announcement"; workbook.description = "All people, titles, notes, policy and channel facts trace to frozen AutomationBench 5112 inputs."
    fixed_dt = datetime(2026, 4, 15, 9, 0, 0); workbook.created = fixed_dt; workbook.modified = fixed_dt

    source = workbook.create_sheet("Source_Roster"); source.append(["Source Row", "Employee", "Department", "Previous Title", "New Title", "Type", "Notes"])
    for item in promotions: source.append([item.source_row, item.employee, item.department, item.previous_title, item.new_title, item.move_type, item.notes])
    add_table(source, "PromotionSourceRoster"); style_sheet(source, {"A":12,"B":23,"C":20,"D":29,"E":31,"F":15,"G":50}, "Frozen Q1 2026 promotion roster · AutomationBench 5112", a3=True)
    for row in range(2,10): source.row_dimensions[row].height = 40

    prior = workbook.create_sheet("Prior_Announcement_Format"); prior.append(["Message ID", "From", "To", "Subject", "Date", "Full Body", "Source Trace"])
    record = messages["msg_last_promo_5112"]; prior.append([record["message_id"],record["from"],record["to"],record["subject"],record["date"],record["body"],"Promotion_Announcement_Context.mbox:msg_last_promo_5112"])
    add_table(prior, "PriorPromotionFormat"); style_sheet(prior, {"A":30,"B":35,"C":39,"D":46,"E":33,"F":105,"G":66}, "Exact prior-quarter announcement format", a3=True); prior.row_dimensions[2].height = 120

    policy = workbook.create_sheet("Policy_Rules"); policy.append(["Rule Order", "Rule", "Exact Requirement", "Source Trace"])
    for order, rule, requirement in policy_rows(): policy.append([order, rule, requirement, "Promotion_Announcement_Context.mbox:msg_announce_policy_5112"])
    add_table(policy, "PromotionPolicyRules"); style_sheet(policy, {"A":14,"B":30,"C":105,"D":66}, "Four VP People announcement rules · latest authority", a3=True)
    for row in range(2,6): policy.row_dimensions[row].height = 53

    evidence = workbook.create_sheet("Evidence_Register"); evidence.append(["Evidence ID", "Type", "Authority / Subject", "Captured Value / Full Body", "Source Trace"])
    for item in promotions:
        value = f"{item.employee} | {item.department} | {item.previous_title} | {item.new_title} | {item.move_type} | {item.notes}"
        evidence.append([f"E-ROW-{item.source_row}","Source roster row",item.employee,value,f"Q1_Promotion_Roster.xlsx:Approved Promotions row {item.source_row}"])
    for message_id in ("msg_last_promo_5112","msg_announce_policy_5112"):
        record = messages[message_id]; evidence.append([f"E-{message_id}","Authoritative email",f"{record['from']} | {record['subject']}",record["body"],f"Promotion_Announcement_Context.mbox:{message_id}"])
    evidence.append(["E-TIME","Frozen runtime","Workflow context",frozen_time,"Workflow_Context.json:meta.current_time"])
    channel = channels["channels"][0]; evidence.append(["E-CHANNEL-C_GEN_5112","Channel snapshot","C_GEN_5112 / general",json.dumps(channel,ensure_ascii=False,separators=(",",":")),"Slack_Channels.json:C_GEN_5112"])
    add_table(evidence, "PromotionEvidenceRegister"); style_sheet(evidence, {"A":34,"B":24,"C":58,"D":112,"E":68}, "Twelve frozen evidence records · no hidden manager identity", a3=True)
    for row in range(2,14): evidence.row_dimensions[row].height = 47
    evidence.row_dimensions[10].height = 100; evidence.row_dimensions[11].height = 130

    decision = workbook.create_sheet("Eligibility_Decisions"); decision.append(["Source Row","Employee","Department","Previous Title","New Title","Type","Notes","Vertical","Lateral","Hold/Pending","Company Eligible","Exclusion Reason","Company Email Route","Receiving Manager Route","Manager Identity","Manager Draft","Slack Route","Final Status","Source Trace"])
    for row in range(2,10):
        decision.append([
            f"=Source_Roster!A{row}",f"=Source_Roster!B{row}",f"=Source_Roster!C{row}",f"=Source_Roster!D{row}",f"=Source_Roster!E{row}",f"=Source_Roster!F{row}",f"=Source_Roster!G{row}",
            f'=IF(F{row}="Vertical","YES","NO")',f'=IF(F{row}="Lateral","YES","NO")',f'=IF(OR(ISNUMBER(SEARCH("hold",G{row})),ISNUMBER(SEARCH("pending",G{row}))),"YES","NO")',
            f'=IF(AND(H{row}="YES",J{row}="NO"),"YES","NO")',f'=IF(K{row}="YES","NONE_ELIGIBLE",IF(I{row}="YES","LATERAL_COMPANY_ANNOUNCEMENT_PROHIBITED","PENDING_HOLD_NOT_CLEARED"))',
            f'=IF(K{row}="YES","YES","NO")',f'=IF(I{row}="YES","YES","NO")',f'=IF(N{row}="YES","UNAVAILABLE_IN_INPUT","N/A")','="NO"','="NO"',
            f'=IF(K{row}="YES","ELIGIBLE_COMPANY_ANNOUNCEMENT",IF(I{row}="YES","RECEIVING_MANAGER_COMMUNICATION_REQUIRED","PENDING_HOLD_NOT_CLEARED"))',
            f'="Q1_Promotion_Roster.xlsx:Approved Promotions row "&A{row}&" | Promotion_Announcement_Context.mbox:msg_announce_policy_5112 | Promotion_Announcement_Context.mbox:msg_last_promo_5112"',
        ])
    add_table(decision, "PromotionEligibilityDecisions")
    widths = {chr(64+i):width for i,width in enumerate([10,20,18,27,29,13,42,12,12,15,16,48,18,22,25,16,14,44,92],start=1)}
    style_sheet(decision,widths,"Formula-derived eligibility · five announced / three excluded",a3=True)
    for row in range(2,10): decision.row_dimensions[row].height = 58
    decision.conditional_formatting.add("R2:R9",FormulaRule(formula=['R2="ELIGIBLE_COMPANY_ANNOUNCEMENT"'],fill=PatternFill("solid",fgColor=GREEN)))
    decision.conditional_formatting.add("R2:R9",FormulaRule(formula=['R2="RECEIVING_MANAGER_COMMUNICATION_REQUIRED"'],fill=PatternFill("solid",fgColor=AMBER)))
    decision.conditional_formatting.add("R2:R9",FormulaRule(formula=['R2="PENDING_HOLD_NOT_CLEARED"'],fill=PatternFill("solid",fgColor=RED)))

    company = workbook.create_sheet("Company_Announcement_Queue"); company.append(["Announcement Order","Source Row","Employee","Department","Previous Title","New Title","Recipient","Subject","State","Sent","Source Trace"])
    for output_row, decision_row in enumerate((2,3,6,8,9),start=2):
        company.append([output_row-1,f"=Eligibility_Decisions!A{decision_row}",f"=Eligibility_Decisions!B{decision_row}",f"=Eligibility_Decisions!C{decision_row}",f"=Eligibility_Decisions!D{decision_row}",f"=Eligibility_Decisions!E{decision_row}",'="all-company@company.example.com"',f'="{SUBJECT}"','="DRAFT_NOT_SENT"',"=FALSE()",f'="Q1_Promotion_Roster.xlsx:Approved Promotions row "&B{output_row}&" | Promotion_Announcement_Context.mbox:msg_announce_policy_5112"'])
    add_table(company,"CompanyPromotionQueue"); style_sheet(company,{"A":17,"B":12,"C":22,"D":20,"E":28,"F":31,"G":40,"H":44,"I":20,"J":10,"K":86},"Five-person company email queue · zero sent",a3=True)
    for row in range(2,7): company.row_dimensions[row].height = 48

    exclusion = workbook.create_sheet("Exclusion_Queue"); exclusion.append(["Exclusion Order","Source Row","Employee","Type","Previous Title","New Title","Notes","Reason","Follow-up Owner","Manager Identity","Draft Created","Sent","Source Trace"])
    for output_row, decision_row in enumerate((4,5,7),start=2):
        exclusion.append([output_row-1,f"=Eligibility_Decisions!A{decision_row}",f"=Eligibility_Decisions!B{decision_row}",f"=Eligibility_Decisions!F{decision_row}",f"=Eligibility_Decisions!D{decision_row}",f"=Eligibility_Decisions!E{decision_row}",f"=Eligibility_Decisions!G{decision_row}",f"=Eligibility_Decisions!L{decision_row}",f'=IF(Eligibility_Decisions!I{decision_row}="YES","receiving_team_manager","HR_APPROVAL_CLEARANCE")',f"=Eligibility_Decisions!O{decision_row}","=FALSE()","=FALSE()",f'="Q1_Promotion_Roster.xlsx:Approved Promotions row "&B{output_row}&" | Promotion_Announcement_Context.mbox:msg_announce_policy_5112"'])
    add_table(exclusion,"PromotionExclusionQueue"); style_sheet(exclusion,{"A":17,"B":12,"C":22,"D":13,"E":28,"F":29,"G":48,"H":49,"I":31,"J":25,"K":16,"L":10,"M":87},"Three policy exclusions · zero invented drafts",a3=True)
    for row in range(2,5): exclusion.row_dimensions[row].height = 58

    summary = workbook.create_sheet("Control_Summary"); summary.append(["Metric","Formula Value","Expected","Control Interpretation","Source Trace"])
    summary_rows = [
        ("Total roster",'=COUNTA(Eligibility_Decisions!A2:A9)',8,"All source rows covered"),("Vertical records",'=COUNTIF(Eligibility_Decisions!H2:H9,"YES")',6,"Six vertical records"),("Lateral records",'=COUNTIF(Eligibility_Decisions!I2:I9,"YES")',2,"Clara and Faisal"),("Hold/pending records",'=COUNTIF(Eligibility_Decisions!J2:J9,"YES")',1,"Derek only"),("Company eligible",'=COUNTIF(Eligibility_Decisions!K2:K9,"YES")',5,"Ava, Brandon, Elena, Greta, Hugo"),("Excluded",'=COUNTIF(Eligibility_Decisions!K2:K9,"NO")',3,"Two lateral plus Derek"),("Manager routes required",'=COUNTIF(Eligibility_Decisions!N2:N9,"YES")',2,"Lateral moves only"),("Manager identities available",'=COUNTIF(Eligibility_Decisions!O2:O9,"<>UNAVAILABLE_IN_INPUT")-COUNTIF(Eligibility_Decisions!O2:O9,"N/A")',0,"Do not invent identity"),("Company email drafts",'=1',1,"One five-person draft"),("Company emails sent",'=COUNTIF(Company_Announcement_Queue!J2:J6,TRUE())',0,"DRAFT_NOT_SENT"),("Manager notification drafts",'=COUNTIF(Exclusion_Queue!K2:K4,TRUE())',0,"Blocked without identity"),("Slack drafts",'=0',0,"Channel not authorized"),("Slack posts",'=0',0,"No side effects")]
    for metric,formula,expected,interpretation in summary_rows: summary.append([metric,formula,expected,interpretation,"Eligibility_Decisions + controlled queues"])
    add_table(summary,"PromotionControlSummary"); style_sheet(summary,{"A":37,"B":19,"C":14,"D":57,"E":55},"Formula summary · five eligible / three excluded / zero side effects",a3=True)
    for row in range(2,15): summary.row_dimensions[row].height = 35
    chart = BarChart(); chart.title = "Promotion governance scope"; chart.y_axis.title = "Employees"; chart.x_axis.title = "Decision"; chart.height = 7.3; chart.width = 14
    chart.add_data(Reference(summary,min_col=2,min_row=6,max_row=7),titles_from_data=False); chart.set_categories(Reference(summary,min_col=1,min_row=6,max_row=7)); chart.legend=None; summary.add_chart(chart,"G2"); summary.print_area="A1:N14"

    workbook.calculation.fullCalcOnLoad=True; workbook.calculation.forceFullCalc=True; workbook.calculation.calcMode="auto"
    workbook.save(path); recalculate_xlsx(path); normalize_ooxml(path)


def set_cell_shading(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr(); shading = properties.find(qn("w:shd"))
    if shading is None: shading = OxmlElement("w:shd"); properties.append(shading)
    shading.set(qn("w:fill"), color)


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr(); repeat = OxmlElement("w:tblHeader"); repeat.set(qn("w:val"), "true"); properties.append(repeat)


def format_doc_table(table, widths: list[float] | None = None, font_size: float = 7.2) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.style = "Table Grid"; set_repeat_table_header(table.rows[0])
    for index, row in enumerate(table.rows):
        for column, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths and column < len(widths): cell.width = Cm(widths[column])
            if index == 0: set_cell_shading(cell, NAVY)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Aptos"; run.font.size = Pt(font_size)
                    if index == 0: run.font.bold = True; run.font.color.rgb = RGBColor(255,255,255)


def add_title(document: Document, title: str, subtitle: str) -> None:
    paragraph=document.add_paragraph(); paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER; paragraph.paragraph_format.space_after=Pt(2)
    run=paragraph.add_run(title); run.font.name="Aptos Display"; run.font.size=Pt(17); run.font.bold=True; run.font.color.rgb=RGBColor(23,54,93)
    sub=document.add_paragraph(); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER; sub.paragraph_format.space_after=Pt(5)
    run=sub.add_run(subtitle); run.font.name="Aptos"; run.font.size=Pt(8.5); run.font.color.rgb=RGBColor(89,89,89)


def add_heading(document: Document, text: str) -> None:
    paragraph=document.add_paragraph(); paragraph.paragraph_format.space_before=Pt(4); paragraph.paragraph_format.space_after=Pt(2)
    run=paragraph.add_run(text); run.font.name="Aptos Display"; run.font.size=Pt(11); run.font.bold=True; run.font.color.rgb=RGBColor(23,54,93)


def add_body(document: Document, text: str) -> None:
    paragraph=document.add_paragraph(); paragraph.paragraph_format.space_after=Pt(2); paragraph.paragraph_format.line_spacing=0.92
    run=paragraph.add_run(text); run.font.name="Aptos"; run.font.size=Pt(8)


def write_brief(path: Path, decisions: list[Decision]) -> None:
    document=Document(); section=document.sections[0]; section.orientation=WD_ORIENT.LANDSCAPE; section.page_width,section.page_height=section.page_height,section.page_width
    section.top_margin=Cm(1.0); section.bottom_margin=Cm(0.9); section.left_margin=Cm(1.1); section.right_margin=Cm(1.1)
    document.styles["Normal"].font.name="Aptos"; document.styles["Normal"].font.size=Pt(8)
    document.core_properties.author="EvoCowork AutomationBench 5112 adaptation"; document.core_properties.title="Promotion Announcement Control Brief"; document.core_properties.created=datetime(2026,4,15,9,0,0); document.core_properties.modified=datetime(2026,4,15,9,0,0)

    add_title(document,"Q1 Promotion Announcement Control Brief","AutomationBench 5112 · policy-governed company announcement · draft only")
    add_heading(document,"Executive result")
    add_body(document,"Eight roster records resolve to five ELIGIBLE_COMPANY_ANNOUNCEMENT and three policy exclusions. The company announcement includes Ava Chen, Brandon Osei, Elena Vasquez, Greta Nilsson and Hugo Fernandez in source order. Clara Dubois and Faisal Rahman are lateral moves assigned to receiving-team manager communication, but no manager identity exists in the inputs, so no manager draft is created. Derek Holm is withheld because his pending note is not cleared. Company email drafts=1; company emails sent=0; manager drafts=0; Slack drafts=0; Slack posts=0; real side effects=0.")
    table=document.add_table(rows=1,cols=5)
    for i,text in enumerate(["Decision","Count","People","Route","Execution"]): table.rows[0].cells[i].text=text
    for values in [
        ("Company eligible","5","Ava; Brandon; Elena; Greta; Hugo","all-company email","1 draft / 0 sent"),
        ("Lateral excluded","2","Clara; Faisal","Receiving manager","0 identity / 0 draft"),
        ("Pending hold","1","Derek","Await clearance","0 announcement"),
        ("Slack","0","No names or titles","NOT_AUTHORIZED_CHANNEL","0 draft / 0 post"),
    ]:
        cells=table.add_row().cells
        for i,value in enumerate(values): cells[i].text=value
    format_doc_table(table,[5.0,2.3,8.1,6.4,5.3])
    add_heading(document,"Authority and conflict resolution")
    add_body(document,"The user asked to include lateral moves, but the latest VP People guideline explicitly prohibits lateral moves in company-wide promotion announcements and assigns them to the receiving team's manager. The same guideline blocks hold/pending records and specifies all-company@company.example.com. The authoritative policy therefore overrides the conflicting user instruction. The prior-quarter email supplies presentation format only; it does not override current scope policy.")

    document.add_page_break(); add_title(document,"Complete eight-person eligibility register","Every source row retains department, titles, type, note, decision and route")
    table=document.add_table(rows=1,cols=11)
    for i,text in enumerate(["Row","Employee","Department","Previous title","New title","Type","Hold / pending","Eligible","Company","Manager","Final status"]): table.rows[0].cells[i].text=text
    for item in decisions:
        source=item.source; values=[str(source.source_row),source.employee,source.department,source.previous_title,source.new_title,source.move_type,item.hold_pending,item.company_eligible,item.company_email_route,item.receiving_manager_route,item.final_status]
        cells=table.add_row().cells
        for i,value in enumerate(values): cells[i].text=value
    format_doc_table(table,[1.3,3.0,2.4,4.0,4.3,1.8,2.1,1.8,1.8,1.8,5.5],font_size=6.25)
    add_heading(document,"Company announcement lines")
    queue=document.add_table(rows=1,cols=5)
    for i,text in enumerate(["Order","Employee","Previous title","New title","Department"]): queue.rows[0].cells[i].text=text
    for order,item in enumerate((value for value in decisions if value.company_eligible=="YES"),start=1):
        source=item.source; cells=queue.add_row().cells
        for i,value in enumerate([str(order),source.employee,source.previous_title,source.new_title,source.department]): cells[i].text=value
    format_doc_table(queue,[1.7,4.4,6.5,7.0,4.1])

    document.add_page_break(); add_title(document,"Exclusion, routing and execution controls","Three exclusions · no invented managers · zero external side effects")
    exclusions=document.add_table(rows=1,cols=7)
    for i,text in enumerate(["Employee","Type","Notes","Reason","Follow-up owner","Manager identity","Draft / sent"]): exclusions.rows[0].cells[i].text=text
    rows=[
        ("Clara Dubois","Lateral","Moving from Growth to Platform team","LATERAL_COMPANY_ANNOUNCEMENT_PROHIBITED","receiving_team_manager","UNAVAILABLE_IN_INPUT","0 / 0"),
        ("Derek Holm","Vertical","Pending - awaiting final sign-off from CTO","PENDING_HOLD_NOT_CLEARED","HR_APPROVAL_CLEARANCE","N/A","0 / 0"),
        ("Faisal Rahman","Lateral","Transfer to EMEA team","LATERAL_COMPANY_ANNOUNCEMENT_PROHIBITED","receiving_team_manager","UNAVAILABLE_IN_INPUT","0 / 0"),
    ]
    for values in rows:
        cells=exclusions.add_row().cells
        for i,value in enumerate(values): cells[i].text=value
    format_doc_table(exclusions,[3.3,2.0,7.0,7.2,4.6,4.3,2.0],font_size=6.5)
    add_heading(document,"Frozen policy rules")
    policy=document.add_table(rows=1,cols=3)
    for i,text in enumerate(["Order","Rule","Exact requirement"]): policy.rows[0].cells[i].text=text
    for order,rule,requirement in policy_rows():
        cells=policy.add_row().cells
        for i,value in enumerate([str(order),rule,requirement]): cells[i].text=value
    format_doc_table(policy,[1.5,5.1,20.5],font_size=6.5)
    add_heading(document,"Reconciliation statement")
    add_body(document,"Cross-file totals: source roster=8; vertical=6; lateral=2; pending/hold=1; company eligible=5; excluded=3; company email drafts=1; company emails sent=0; receiving-manager routes required=2; manager identities available=0; manager drafts=0; Slack drafts=0; Slack posts=0; real side effects=0. Both authoritative emails, the channel snapshot and frozen runtime are represented in the audit workbook.")
    document.save(path); normalize_ooxml(path)


def add_textbox(slide,x:float,y:float,w:float,h:float,text:str,*,font_size:float=16,bold:bool=False,color=PPT_NAVY,align=PP_ALIGN.LEFT,fill=None,line=None,radius:bool=False):
    shape_type=MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE; shape=slide.shapes.add_shape(shape_type,Inches(x),Inches(y),Inches(w),Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb=fill if fill is not None else PPT_WHITE; shape.line.color.rgb=line if line is not None else shape.fill.fore_color.rgb
    frame=shape.text_frame; frame.clear(); frame.vertical_anchor=MSO_ANCHOR.MIDDLE; frame.margin_left=Inches(0.08); frame.margin_right=Inches(0.08); frame.margin_top=Inches(0.03); frame.margin_bottom=Inches(0.03)
    paragraph=frame.paragraphs[0]; paragraph.alignment=align; run=paragraph.add_run(); run.text=text; run.font.name="Aptos"; run.font.size=PptPt(font_size); run.font.bold=bold; run.font.color.rgb=color
    return shape


def add_slide_title(slide,title:str,subtitle:str)->None:
    add_textbox(slide,0.55,0.25,12.25,0.55,title,font_size=24,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER)
    add_textbox(slide,0.75,0.78,11.85,0.36,subtitle,font_size=10.5,color=PptRGBColor(89,89,89),align=PP_ALIGN.CENTER)


def style_ppt_table(table,*,header_fill=PPT_NAVY,body_font=9.5)->None:
    for row_index,row in enumerate(table.rows):
        for cell in row.cells:
            cell.fill.solid(); cell.fill.fore_color.rgb=header_fill if row_index==0 else PPT_WHITE; cell.margin_left=Inches(0.05); cell.margin_right=Inches(0.05); cell.margin_top=Inches(0.025); cell.margin_bottom=Inches(0.025)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment=PP_ALIGN.LEFT
                for run in paragraph.runs: run.font.name="Aptos"; run.font.size=PptPt(body_font); run.font.bold=row_index==0; run.font.color.rgb=PPT_WHITE if row_index==0 else PPT_NAVY


def write_presentation(path:Path,decisions:list[Decision])->None:
    presentation=Presentation(); presentation.slide_width=Inches(13.333); presentation.slide_height=Inches(7.5)
    presentation.core_properties.author="EvoCowork AutomationBench 5112 adaptation"; presentation.core_properties.title="Promotion Announcement Review"; presentation.core_properties.created=datetime(2026,4,15,9,0,0); presentation.core_properties.modified=datetime(2026,4,15,9,0,0)
    blank=presentation.slide_layouts[6]

    slide=presentation.slides.add_slide(blank); add_slide_title(slide,"Q1 promotion announcement governance","Eight reviewed · five company eligible · three excluded · zero real side effects")
    for x,value,label,fill in [(0.65,"8","roster records",PPT_LIGHT_BLUE),(3.15,"5","company eligible",PPT_GREEN),(5.65,"2","lateral excluded",PPT_AMBER),(8.15,"1","pending hold",PPT_RED),(10.65,"0","real side effects",PPT_AMBER)]:
        add_textbox(slide,x,1.25,2.05,0.72,value,font_size=25,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=fill,line=PPT_BLUE,radius=True); add_textbox(slide,x,1.92,2.05,0.34,label,font_size=9.5,color=PPT_NAVY,align=PP_ALIGN.CENTER)
    table_shape=slide.shapes.add_table(6,5,Inches(0.65),Inches(2.5),Inches(12.05),Inches(3.55)); table=table_shape.table
    for i,text in enumerate(["Order","Employee","Previous title","New title","Department"]): table.cell(0,i).text=text
    eligible=[item for item in decisions if item.company_eligible=="YES"]
    for row,item in enumerate(eligible,start=1):
        source=item.source
        for col,value in enumerate([str(row),source.employee,source.previous_title,source.new_title,source.department]): table.cell(row,col).text=value
    for i,width in enumerate([0.8,2.4,3.1,3.4,2.35]): table.columns[i].width=Inches(width)
    style_ppt_table(table,body_font=10.1)
    add_textbox(slide,0.65,6.3,12.05,0.47,"One all-company draft · email sent=0 · manager drafts=0 · Slack drafts/posts=0",font_size=10.5,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_AMBER,line=PPT_AMBER,radius=True)

    slide=presentation.slides.add_slide(blank); add_slide_title(slide,"Complete eight-person eligibility register","Greta is included by full policy evaluation even though upstream assertions do not name her")
    table_shape=slide.shapes.add_table(9,9,Inches(0.35),Inches(1.25),Inches(12.63),Inches(5.15)); table=table_shape.table
    for i,text in enumerate(["Row","Employee","Department","Type","Hold","Eligible","Company","Manager","Final status"]): table.cell(0,i).text=text
    for row,item in enumerate(decisions,start=1):
        source=item.source; values=[str(source.source_row),source.employee,source.department,source.move_type,item.hold_pending,item.company_eligible,item.company_email_route,item.receiving_manager_route,item.final_status]
        for col,value in enumerate(values): table.cell(row,col).text=value
    for i,width in enumerate([0.6,1.9,1.45,1.05,0.7,0.85,0.85,0.85,4.38]): table.columns[i].width=Inches(width)
    style_ppt_table(table,body_font=8.1)
    add_textbox(slide,0.65,6.47,12.03,0.42,"Policy wins over the request to include lateral moves · no manager identity is invented",font_size=10.5,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_LIGHT_BLUE,line=PPT_LIGHT_BLUE,radius=True)

    slide=presentation.slides.add_slide(blank); add_slide_title(slide,"Exclusion and channel controls","Latest VP People policy governs scope; prior-quarter email governs format only")
    exclusions=[item for item in decisions if item.company_eligible=="NO"]
    table_shape=slide.shapes.add_table(4,6,Inches(0.45),Inches(1.25),Inches(12.43),Inches(2.65)); table=table_shape.table
    for i,text in enumerate(["Employee","Type","Notes","Reason","Owner","Draft / sent"]): table.cell(0,i).text=text
    for row,item in enumerate(exclusions,start=1):
        source=item.source; owner="receiving_team_manager" if item.lateral=="YES" else "HR_APPROVAL_CLEARANCE"; values=[source.employee,source.move_type,source.notes,item.exclusion_reason,owner,"0 / 0"]
        for col,value in enumerate(values): table.cell(row,col).text=value
    for i,width in enumerate([1.65,0.9,3.3,3.6,2.15,0.83]): table.columns[i].width=Inches(width)
    style_ppt_table(table,body_font=8.5)
    add_textbox(slide,0.45,4.25,3.85,1.08,"Company email\n1 draft · 0 sent\nall-company@company.example.com",font_size=13,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_GREEN,line=PPT_BLUE,radius=True)
    add_textbox(slide,4.74,4.25,3.85,1.08,"Receiving managers\n2 routes required\n0 identities · 0 drafts",font_size=13,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_AMBER,line=PPT_BLUE,radius=True)
    add_textbox(slide,9.03,4.25,3.85,1.08,"#general Slack\nNOT_AUTHORIZED_CHANNEL\n0 drafts · 0 posts",font_size=13,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_AMBER,line=PPT_BLUE,radius=True)
    add_textbox(slide,0.45,5.75,12.43,0.62,"DRAFT_NOT_SENT · BLOCKED_IDENTITY_UNAVAILABLE · no excluded employee in company email or Slack",font_size=11,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_LIGHT_BLUE,line=PPT_LIGHT_BLUE,radius=True)
    add_textbox(slide,0.45,6.5,12.43,0.4,"Company emails sent=0 · Manager notifications sent=0 · Slack posts=0 · Real side effects=0",font_size=10.5,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_GREEN,line=PPT_GREEN,radius=True)

    presentation.save(path); normalize_ooxml(path)


def announcement_body(decisions:list[Decision])->str:
    lines=["Please join us in congratulating the following team members on their well-deserved promotions!",""]
    for item in decisions:
        if item.company_eligible=="YES":
            source=item.source; lines.append(f"- {source.employee}: {source.previous_title} -> {source.new_title} ({source.department})")
    lines.extend(["","Congratulations to all!","- HR Team","","DRAFT_NOT_SENT — this offline benchmark artifact has not been sent."])
    return "\n".join(lines)


def write_mbox(path:Path,decisions:list[Decision])->None:
    message=EmailMessage(); message["From"]="hr-ops@company.example.com"; message["To"]="all-company@company.example.com"; message["Subject"]=SUBJECT; message["Date"]="Wed, 15 Apr 2026 09:00:00 +0000"; message["Message-ID"]="<promotion-q1-2026-5112@evocowork.local>"; message["X-EvoCowork-Status"]="DRAFT_NOT_SENT"
    message.set_content(announcement_body(decisions),charset="utf-8",cte="quoted-printable"); stream=BytesIO(); BytesGenerator(stream,policy=SMTP.clone(linesep="\n",max_line_length=78)).flatten(message,unixfrom=False); block=stream.getvalue()
    with path.open("wb") as handle:
        handle.write(b"From MAILER-DAEMON Wed Apr 15 09:00:00 2026\n"); handle.write(block.replace(b"\nFrom ",b"\n>From "))
        if not block.endswith(b"\n"): handle.write(b"\n")
        handle.write(b"\n")


def write_exclusions(path:Path,decisions:list[Decision])->None:
    excluded=[]
    for item in decisions:
        if item.company_eligible=="YES": continue
        source=item.source; lateral=item.lateral=="YES"
        excluded.append({"source_row":source.source_row,"employee":source.employee,"department":source.department,"previous_title":source.previous_title,"new_title":source.new_title,"type":source.move_type,"notes":source.notes,"exclusion_reason":item.exclusion_reason,"final_status":item.final_status,"follow_up_owner":"receiving_team_manager" if lateral else "HR_APPROVAL_CLEARANCE","manager_identity":"UNAVAILABLE_IN_INPUT" if lateral else "N/A","draft_created":False,"sent":False,"source_trace":[f"Q1_Promotion_Roster.xlsx:Approved Promotions row {source.source_row}","Promotion_Announcement_Context.mbox:msg_announce_policy_5112"]})
    payload={"task_id":"automationbench-5112","status":"SCOPE_CONTROLLED","excluded_count":3,"company_announcement_excluded_count":3,"draft_count":0,"sent_count":0,"excluded":excluded}
    path.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")


def write_slack(path:Path)->None:
    payload={"task_id":"automationbench-5112","source_channel":{"id":"C_GEN_5112","name":"general","is_private":False},"decision":"NOT_AUTHORIZED_CHANNEL","reason":"The latest VP People guideline routes company promotion announcements to all-company@company.example.com and does not authorize Slack.","draft_count":0,"posted_count":0,"posted":False,"employee_names_in_slack":[],"titles_in_slack":[],"messages_created":[],"source_trace":["Promotion_Announcement_Context.mbox:msg_announce_policy_5112","Slack_Channels.json:C_GEN_5112","Promotion_Eligibility_Audit.xlsx:Control_Summary"]}
    path.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")


def write_execution(path:Path)->None:
    payload={"task_id":"automationbench-5112","status":"CONTROLLED_DRAFT_ONLY","company_email":{"recipient":"all-company@company.example.com","subject":SUBJECT,"eligible_count":5,"draft_count":1,"sent_count":0,"state":"DRAFT_NOT_SENT"},"lateral_manager_communications":{"required_count":2,"manager_identity_available_count":0,"draft_count":0,"sent_count":0,"state":"BLOCKED_IDENTITY_UNAVAILABLE"},"pending_hold":{"employee":"Derek Holm","count":1,"announcement_draft_count":0,"state":"PENDING_HOLD_NOT_CLEARED"},"slack":{"channel":"C_GEN_5112 / general","draft_count":0,"posted_count":0,"state":"NOT_AUTHORIZED_CHANNEL"},"real_side_effect_count":0,"source_trace":["Q1_Promotion_Roster.xlsx:Approved Promotions rows 2-9","Promotion_Announcement_Context.mbox:msg_announce_policy_5112","Promotion_Announcement_Context.mbox:msg_last_promo_5112","Slack_Channels.json:C_GEN_5112"]}
    path.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")


def main()->None:
    parser=argparse.ArgumentParser(); parser.add_argument("--input-dir",type=Path,required=True); parser.add_argument("--output-dir",type=Path,required=True); args=parser.parse_args()
    promotions,messages,channels,frozen_time=read_inputs(args.input_dir); decisions=derive_decisions(promotions)
    args.output_dir.mkdir(parents=True,exist_ok=True)
    for item in args.output_dir.iterdir():
        if item.is_file(): item.unlink()
        elif item.is_dir(): shutil.rmtree(item)
    write_workbook(args.output_dir/XLSX_NAME,promotions,messages,channels,frozen_time); write_brief(args.output_dir/DOCX_NAME,decisions); write_presentation(args.output_dir/PPTX_NAME,decisions); write_mbox(args.output_dir/MBOX_NAME,decisions); write_exclusions(args.output_dir/EXCLUSION_NAME,decisions); write_slack(args.output_dir/SLACK_NAME); write_execution(args.output_dir/EXECUTION_NAME)
    expected={XLSX_NAME,DOCX_NAME,PPTX_NAME,MBOX_NAME,EXCLUSION_NAME,SLACK_NAME,EXECUTION_NAME}
    if {item.name for item in args.output_dir.iterdir() if item.is_file()}!=expected: raise RuntimeError("交付文件集合与冻结任务不一致")


if __name__=="__main__": main()
