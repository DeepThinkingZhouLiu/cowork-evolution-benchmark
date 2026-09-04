#!/usr/bin/env python3
"""向 AutomationBench 5112 Golden 逐项注入本题专属错误，验证 Judge 拒绝坏答案。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
import zipfile
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path
from typing import Callable

from docx import Document
from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
XLSX = "Promotion_Eligibility_Audit.xlsx"
DOCX = "Promotion_Announcement_Control_Brief.docx"
PPTX = "Promotion_Announcement_Review.pptx"
MBOX = "Company_Promotion_Announcement_Draft.mbox"
EXCLUSION = "Promotion_Exclusion_Register.json"
SLACK = "Promotion_Slack_Control.json"
EXECUTION = "Promotion_Execution_Control.json"
ELIGIBLE = ["Ava Chen", "Brandon Osei", "Elena Vasquez", "Greta Nilsson", "Hugo Fernandez"]
EXCLUDED = ["Clara Dubois", "Derek Holm", "Faisal Rahman"]
Mutator = Callable[[Path], None]
Case = tuple[str, Mutator, str]


def load_judge():
    spec = importlib.util.spec_from_file_location("promotion_5112_judge", ROOT / "tests" / "judge.py")
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入本题 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    module.render = lambda path: ({".xlsx": 8, ".docx": 3, ".pptx": 3}[path.suffix], "render-ok")
    return module


def rewrite_zip(path: Path, transform: Callable[[str, bytes], bytes | None]) -> None:
    temporary = path.with_suffix(path.suffix + ".negative")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = transform(info.filename, source.read(info.filename))
            if data is not None:
                target.writestr(info, data)
    temporary.replace(path)


def sheet_xml_path(archive: zipfile.ZipFile, sheet_name: str) -> str:
    workbook = etree.fromstring(archive.read("xl/workbook.xml"))
    relationships = etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    namespaces = {
        "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    relationship_id = workbook.xpath(f'//m:sheet[@name="{sheet_name}"]/@r:id', namespaces=namespaces)[0]
    target = relationships.xpath(f'//pr:Relationship[@Id="{relationship_id}"]/@Target', namespaces=namespaces)[0]
    return target.lstrip("/") if target.startswith("/") else "xl/" + target.lstrip("/")


def edit_cell(path: Path, sheet_name: str, coordinate: str, *, value=None, remove_formula: bool = False) -> None:
    with zipfile.ZipFile(path, "r") as archive:
        target_sheet = sheet_xml_path(archive, sheet_name)
    changed = False
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

    def transform(name: str, data: bytes) -> bytes:
        nonlocal changed
        if name != target_sheet:
            return data
        root = etree.fromstring(data)
        cells = root.xpath(f'//m:c[@r="{coordinate}"]', namespaces={"m": namespace})
        if not cells:
            raise ValueError(f"{sheet_name}!{coordinate} 不存在")
        cell = cells[0]
        if remove_formula:
            formula = cell.find(f"{{{namespace}}}f")
            if formula is None:
                raise ValueError(f"{sheet_name}!{coordinate} 没有公式")
            cell.remove(formula)
        if value is not None:
            node = cell.find(f"{{{namespace}}}v")
            if node is None:
                node = etree.SubElement(cell, f"{{{namespace}}}v")
            if isinstance(value, bool):
                node.text = "1" if value else "0"
                cell.set("t", "b")
            elif isinstance(value, int):
                node.text = str(value)
                cell.attrib.pop("t", None)
            else:
                node.text = str(value)
                cell.set("t", "str")
        changed = True
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    rewrite_zip(path, transform)
    if not changed:
        raise ValueError("目标单元格未修改")


def openpyxl_mutation(path: Path, mutator: Callable) -> None:
    workbook = load_workbook(path)
    mutator(workbook)
    workbook.save(path)


def mutate_json(path: Path, mutator: Callable[[dict], None]) -> None:
    payload = json.loads(path.read_text(encoding="utf-8"))
    mutator(payload)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def mbox_states(path: Path) -> list[dict]:
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    states = []
    for message in messages:
        payload = message.get_payload(decode=True) or b""
        states.append({
            "from": str(message.get("From", "")), "to": str(message.get("To", "")),
            "cc": str(message.get("Cc", "")), "bcc": str(message.get("Bcc", "")),
            "subject": str(message.get("Subject", "")), "date": str(message.get("Date", "")),
            "message_id": str(message.get("Message-ID", "")),
            "status": str(message.get("X-EvoCowork-Status", "")),
            "body": payload.decode(message.get_content_charset() or "utf-8", errors="replace"),
        })
    return states


def write_mbox_states(path: Path, states: list[dict]) -> None:
    chunks: list[bytes] = []
    for state in states:
        message = EmailMessage(policy=SMTP)
        for key, header in (("from", "From"), ("to", "To"), ("cc", "Cc"), ("bcc", "Bcc"), ("subject", "Subject"), ("date", "Date"), ("message_id", "Message-ID"), ("status", "X-EvoCowork-Status")):
            if state.get(key):
                message[header] = state[key]
        message.set_content(state["body"].rstrip() + "\n", charset="utf-8")
        buffer = BytesIO()
        BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
        chunks.extend([b"From MAILER-DAEMON Wed Apr 15 09:00:00 2026\n", buffer.getvalue().replace(b"\r\n", b"\n"), b"\n"])
    path.write_bytes(b"".join(chunks))


def mutate_message(path: Path, mutator: Callable[[dict], None]) -> None:
    states = mbox_states(path)
    if len(states) != 1:
        raise ValueError("预期 Golden 中恰好一封邮件")
    mutator(states[0])
    write_mbox_states(path, states)


def body_replace(state: dict, old: str, new: str) -> None:
    if old not in state["body"]:
        raise ValueError(f"邮件正文找不到 {old!r}")
    state["body"] = state["body"].replace(old, new, 1)


def replace_docx_text(path: Path, old: str, new: str, *, row_token: str | None = None) -> None:
    document = Document(path)
    changed = False
    if row_token is None:
        for paragraph in document.paragraphs:
            if old in paragraph.text:
                paragraph.text = paragraph.text.replace(old, new, 1)
                changed = True
                break
    if not changed:
        for table in document.tables:
            for row in table.rows:
                if row_token and row_token not in " | ".join(cell.text for cell in row.cells):
                    continue
                for cell in row.cells:
                    if old in cell.text:
                        cell.text = cell.text.replace(old, new, 1)
                        changed = True
                        break
                if changed:
                    break
            if changed:
                break
    if not changed:
        raise ValueError(f"Word 找不到 {old!r}")
    document.save(path)


def replace_pptx_text(path: Path, old: str, new: str) -> None:
    presentation = Presentation(path)
    changed = False
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and old in shape.text:
                shape.text = shape.text.replace(old, new, 1)
                changed = True
                break
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        if old in cell.text:
                            cell.text = cell.text.replace(old, new, 1)
                            changed = True
                            break
                    if changed:
                        break
            if changed:
                break
        if changed:
            break
    if not changed:
        raise ValueError(f"PPT 找不到 {old!r}")
    presentation.save(path)


def drop_ppt_slide(path: Path) -> None:
    presentation = Presentation(path)
    presentation.slides._sldIdLst.remove(presentation.slides._sldIdLst[-1])
    presentation.save(path)


def remove_path(root: Path, name: str) -> None:
    (root / name).unlink()


def corrupt_path(root: Path, name: str) -> None:
    (root / name).write_bytes(b"not-a-valid-artifact")


def build_cases() -> list[Case]:
    cases: list[Case] = []

    for name in (XLSX, DOCX, PPTX, MBOX, EXCLUSION, SLACK, EXECUTION):
        cases.append((f"missing-{name}", lambda root, n=name: remove_path(root, n), "R001"))
    cases.append(("extra-file", lambda root: (root / "unexpected.txt").write_text("unexpected", encoding="utf-8"), "R001"))
    for name in (XLSX, DOCX, PPTX, MBOX, EXCLUSION, SLACK, EXECUTION):
        expected = "R011" if name == MBOX else "R001"
        cases.append((f"corrupt-{name}", lambda root, n=name: corrupt_path(root, n), expected))

    source_edits = [
        ("source-ava-name", "B2", "Ava Wrong"), ("source-brandon-department", "C3", "Finance"),
        ("source-clara-previous", "D4", "Senior Product Manager"), ("source-derek-new", "E5", "Principal Engineer"),
        ("source-elena-type", "F6", "Lateral"), ("source-faisal-notes", "G7", "Transfer to APAC team"),
        ("source-greta-name", "B8", "Greta Wrong"), ("source-hugo-row", "A9", 10),
        ("source-ava-new", "E2", "Staff Software Engineer"), ("source-brandon-previous", "D3", "Sales Manager"),
        ("source-clara-type", "F4", "Vertical"), ("source-derek-notes", "G5", "Approved"),
        ("source-elena-name", "B6", "Elena Wrong"), ("source-faisal-department", "C7", "Sales"),
        ("source-greta-new", "E8", "Support Manager"), ("source-hugo-previous", "D9", "Engineer"),
    ]
    for name, cell, value in source_edits:
        cases.append((name, lambda root, c=cell, v=value: edit_cell(root / XLSX, "Source_Roster", c, value=v), "R002"))

    context_edits = [
        ("prior-message-id", "Prior_Announcement_Format", "A2", "wrong-id"),
        ("prior-from", "Prior_Announcement_Format", "B2", "wrong@company.example.com"),
        ("prior-recipient", "Prior_Announcement_Format", "C2", "leaders@company.example.com"),
        ("prior-subject", "Prior_Announcement_Format", "D2", "Wrong quarter"),
        ("prior-body", "Prior_Announcement_Format", "F2", "No reusable format"),
        ("prior-trace", "Prior_Announcement_Format", "G2", "wrong-source"),
        ("policy-scope", "Policy_Rules", "C2", "Announce all changes"),
        ("policy-lateral", "Policy_Rules", "C3", "Include lateral moves"),
        ("policy-recipient", "Policy_Rules", "C4", "Announcements go to Slack"),
        ("policy-pending", "Policy_Rules", "C5", "Pending employees may be announced"),
        ("policy-order", "Policy_Rules", "A3", 9),
        ("evidence-ava", "Evidence_Register", "D2", "Ava Chen | wrong"),
        ("evidence-clara-note", "Evidence_Register", "D4", "Clara Dubois | Product | Lateral"),
        ("evidence-derek-note", "Evidence_Register", "D5", "Derek Holm | no pending note"),
        ("evidence-faisal-note", "Evidence_Register", "D7", "Faisal Rahman | no team"),
        ("evidence-greta-trace", "Evidence_Register", "E8", "wrong-row"),
        ("evidence-prior", "Evidence_Register", "D10", "wrong prior email"),
        ("evidence-policy", "Evidence_Register", "D11", "wrong policy email"),
        ("evidence-time", "Evidence_Register", "D12", "2026-04-16T09:00:00Z"),
        ("evidence-channel", "Evidence_Register", "D13", '{"id":"C_GEN_5112","name":"general","is_private":true}'),
    ]
    for name, sheet, cell, value in context_edits:
        cases.append((name, lambda root, s=sheet, c=cell, v=value: edit_cell(root / XLSX, s, c, value=v), "R003"))

    formula_cells = [
        ("formula-source-row", "Eligibility_Decisions", "A2"), ("formula-source-name", "Eligibility_Decisions", "B3"),
        ("formula-source-notes", "Eligibility_Decisions", "G5"), ("formula-vertical", "Eligibility_Decisions", "H4"),
        ("formula-lateral", "Eligibility_Decisions", "I7"), ("formula-hold", "Eligibility_Decisions", "J5"),
        ("formula-eligible", "Eligibility_Decisions", "K8"), ("formula-exclusion", "Eligibility_Decisions", "L4"),
        ("formula-company-route", "Eligibility_Decisions", "M9"), ("formula-manager-route", "Eligibility_Decisions", "N7"),
        ("formula-manager-identity", "Eligibility_Decisions", "O4"), ("formula-manager-draft", "Eligibility_Decisions", "P4"),
        ("formula-slack-route", "Eligibility_Decisions", "Q2"), ("formula-final-status", "Eligibility_Decisions", "R5"),
        ("formula-source-trace", "Eligibility_Decisions", "S8"),
        ("formula-company-queue", "Company_Announcement_Queue", "C5"),
        ("formula-company-sent", "Company_Announcement_Queue", "J2"),
        ("formula-exclusion-queue", "Exclusion_Queue", "H2"),
        ("formula-manager-draft-queue", "Exclusion_Queue", "K4"),
        ("formula-summary", "Control_Summary", "B6"),
    ]
    for name, sheet, cell in formula_cells:
        cases.append((name, lambda root, s=sheet, c=cell: edit_cell(root / XLSX, s, c, remove_formula=True), "R004"))

    decision_edits = [
        ("ava-wrong-department", "C2", "Product", "R005"), ("ava-not-eligible", "K2", "NO", "R005"),
        ("ava-no-company-route", "M2", "NO", "R005"), ("ava-wrong-status", "R2", "PENDING_HOLD_NOT_CLEARED", "R005"),
        ("brandon-wrong-title", "E3", "Staff Account Executive", "R006"), ("brandon-lateral", "I3", "YES", "R006"),
        ("brandon-not-eligible", "K3", "NO", "R006"), ("brandon-wrong-row", "A3", 4, "R006"),
        ("elena-wrong-department", "C6", "Product", "R007"), ("elena-hold", "J6", "YES", "R007"),
        ("elena-no-route", "M6", "NO", "R007"), ("elena-wrong-title", "D6", "Senior Tech Lead", "R007"),
        ("greta-omitted", "K8", "NO", "R008"), ("greta-wrong-new-title", "E8", "Support Manager", "R008"),
        ("greta-wrong-status", "R8", "PENDING_HOLD_NOT_CLEARED", "R008"), ("greta-manager-route", "N8", "YES", "R008"),
        ("hugo-wrong-previous", "D9", "Engineer", "R009"), ("hugo-not-eligible", "K9", "NO", "R009"),
        ("hugo-slack-route", "Q9", "YES", "R009"), ("hugo-wrong-row", "A9", 8, "R009"),
        ("clara-company-eligible", "K4", "YES", "R010"), ("clara-manager-invented", "O4", "manager@company.example.com", "R010"),
        ("clara-manager-draft", "P4", "YES", "R010"), ("clara-wrong-status", "R4", "ELIGIBLE_COMPANY_ANNOUNCEMENT", "R010"),
        ("derek-no-hold", "J5", "NO", "R010"), ("derek-company-route", "M5", "YES", "R010"),
        ("derek-wrong-reason", "L5", "NONE_ELIGIBLE", "R010"), ("faisal-manager-invented", "O7", "emea-manager@company.example.com", "R010"),
        ("faisal-company-route", "M7", "YES", "R010"), ("faisal-manager-draft", "P7", "YES", "R010"),
    ]
    for name, cell, value, criterion in decision_edits:
        cases.append((name, lambda root, c=cell, v=value: edit_cell(root / XLSX, "Eligibility_Decisions", c, value=v), criterion))

    queue_edits = [
        ("queue-ava-name", "Company_Announcement_Queue", "C2", "Ava Wrong", "R011"),
        ("queue-brandon-title", "Company_Announcement_Queue", "F3", "Staff Account Executive", "R011"),
        ("queue-elena-row", "Company_Announcement_Queue", "B4", 7, "R011"),
        ("queue-greta-omitted", "Company_Announcement_Queue", "C5", "Derek Holm", "R011"),
        ("queue-hugo-order", "Company_Announcement_Queue", "A6", 4, "R011"),
        ("queue-wrong-recipient", "Company_Announcement_Queue", "G2", "leaders@company.example.com", "R011"),
        ("queue-wrong-subject", "Company_Announcement_Queue", "H2", "Wrong subject", "R011"),
        ("queue-marked-sent", "Company_Announcement_Queue", "J2", True, "R011"),
        ("exclude-clara-reason", "Exclusion_Queue", "H2", "NONE_ELIGIBLE", "R010"),
        ("exclude-clara-manager", "Exclusion_Queue", "J2", "manager@company.example.com", "R010"),
        ("exclude-derek-draft", "Exclusion_Queue", "K3", True, "R010"),
        ("exclude-faisal-sent", "Exclusion_Queue", "L4", True, "R010"),
        ("summary-eligible", "Control_Summary", "B6", 4, "R004"),
        ("summary-sent", "Control_Summary", "B11", 1, "R004"),
    ]
    for name, sheet, cell, value, criterion in queue_edits:
        cases.append((name, lambda root, s=sheet, c=cell, v=value: edit_cell(root / XLSX, s, c, value=v), criterion))

    cases.extend([
        ("xlsx-native-table-removed", lambda root: openpyxl_mutation(root / XLSX, lambda wb: wb["Source_Roster"].tables.clear()), "R016"),
        ("xlsx-native-chart-removed", lambda root: openpyxl_mutation(root / XLSX, lambda wb: setattr(wb["Control_Summary"], "_charts", [])), "R016"),
        ("xlsx-conditional-formatting-removed", lambda root: openpyxl_mutation(root / XLSX, lambda wb: setattr(wb["Eligibility_Decisions"], "conditional_formatting", type(wb["Eligibility_Decisions"].conditional_formatting)())), "R016"),
    ])

    exclusion_cases = [
        ("exclusion-task-id", lambda p: p.update({"task_id": "wrong"})),
        ("exclusion-status", lambda p: p.update({"status": "ANNOUNCED"})),
        ("exclusion-count", lambda p: p.update({"excluded_count": 2})),
        ("exclusion-draft-count", lambda p: p.update({"draft_count": 2})),
        ("exclusion-sent-count", lambda p: p.update({"sent_count": 1})),
        ("exclusion-drop-clara", lambda p: p["excluded"].pop(0)),
        ("exclusion-reverse", lambda p: p["excluded"].reverse()),
        ("exclusion-clara-name", lambda p: p["excluded"][0].update({"employee": "Clara Wrong"})),
        ("exclusion-clara-manager", lambda p: p["excluded"][0].update({"manager_identity": "manager@company.example.com"})),
        ("exclusion-clara-draft", lambda p: p["excluded"][0].update({"draft_created": True})),
        ("exclusion-derek-note", lambda p: p["excluded"][1].update({"notes": "Approved"})),
        ("exclusion-derek-reason", lambda p: p["excluded"][1].update({"exclusion_reason": "NONE_ELIGIBLE"})),
        ("exclusion-derek-sent", lambda p: p["excluded"][1].update({"sent": True})),
        ("exclusion-faisal-owner", lambda p: p["excluded"][2].update({"follow_up_owner": "HR"})),
        ("exclusion-faisal-trace", lambda p: p["excluded"][2].update({"source_trace": []})),
    ]
    for name, mutator in exclusion_cases:
        cases.append((name, lambda root, m=mutator: mutate_json(root / EXCLUSION, m), "R010"))

    cases.extend([
        ("mbox-empty", lambda root: write_mbox_states(root / MBOX, []), "R011"),
        ("mbox-two-messages", lambda root: write_mbox_states(root / MBOX, mbox_states(root / MBOX) * 2), "R011"),
        ("mbox-recipient", lambda root: mutate_message(root / MBOX, lambda s: s.update({"to": "leaders@company.example.com"})), "R011"),
        ("mbox-cc", lambda root: mutate_message(root / MBOX, lambda s: s.update({"cc": "ceo@company.example.com"})), "R011"),
        ("mbox-bcc", lambda root: mutate_message(root / MBOX, lambda s: s.update({"bcc": "audit@company.example.com"})), "R011"),
        ("mbox-from", lambda root: mutate_message(root / MBOX, lambda s: s.update({"from": "vp-people@company.example.com"})), "R011"),
        ("mbox-subject", lambda root: mutate_message(root / MBOX, lambda s: s.update({"subject": "Q1 Moves"})), "R011"),
        ("mbox-sent", lambda root: mutate_message(root / MBOX, lambda s: s.update({"status": "SENT"})), "R011"),
        ("mbox-drop-ava", lambda root: mutate_message(root / MBOX, lambda s: body_replace(s, "- Ava Chen: Software Engineer -> Senior Software Engineer (Engineering)", "")), "R011"),
        ("mbox-drop-greta", lambda root: mutate_message(root / MBOX, lambda s: body_replace(s, "- Greta Nilsson: Support Specialist -> Senior Support Specialist (Support)", "")), "R011"),
        ("mbox-wrong-elena-title", lambda root: mutate_message(root / MBOX, lambda s: body_replace(s, "Tech Lead -> Engineering Manager", "Tech Lead -> Senior Tech Lead")), "R011"),
        ("mbox-reorder", lambda root: mutate_message(root / MBOX, lambda s: body_replace(s, "- Ava Chen: Software Engineer -> Senior Software Engineer (Engineering)\n- Brandon Osei: Account Executive -> Senior Account Executive (Sales)", "- Brandon Osei: Account Executive -> Senior Account Executive (Sales)\n- Ava Chen: Software Engineer -> Senior Software Engineer (Engineering)")), "R011"),
        ("mbox-add-clara", lambda root: mutate_message(root / MBOX, lambda s: s.update({"body": s["body"] + "\n- Clara Dubois: Product Manager -> Product Manager (Product)\n"})), "R011"),
        ("mbox-add-derek", lambda root: mutate_message(root / MBOX, lambda s: s.update({"body": s["body"] + "\n- Derek Holm: Senior Engineer -> Staff Engineer (Engineering)\n"})), "R011"),
        ("mbox-false-sent-body", lambda root: mutate_message(root / MBOX, lambda s: body_replace(s, "has not been sent", "has been sent")), "R011"),
    ])

    slack_cases = [
        ("slack-task-id", lambda p: p.update({"task_id": "wrong"})),
        ("slack-channel-id", lambda p: p["source_channel"].update({"id": "C_OTHER"})),
        ("slack-channel-private", lambda p: p["source_channel"].update({"is_private": True})),
        ("slack-decision-post", lambda p: p.update({"decision": "POST"})),
        ("slack-draft", lambda p: p.update({"draft_count": 1})),
        ("slack-posted-count", lambda p: p.update({"posted_count": 1})),
        ("slack-posted", lambda p: p.update({"posted": True})),
        ("slack-name-leak", lambda p: p["employee_names_in_slack"].append("Greta Nilsson")),
        ("slack-title-leak", lambda p: p["titles_in_slack"].append("Senior Support Specialist")),
        ("slack-message-created", lambda p: p["messages_created"].append({"channel": "general", "text": "Congrats"})),
        ("slack-wrong-reason", lambda p: p.update({"reason": "Slack authorized"})),
        ("slack-remove-trace", lambda p: p.update({"source_trace": []})),
    ]
    for name, mutator in slack_cases:
        cases.append((name, lambda root, m=mutator: mutate_json(root / SLACK, m), "R012"))

    execution_cases = [
        ("execution-task-id", lambda p: p.update({"task_id": "wrong"}), "R012"),
        ("execution-status", lambda p: p.update({"status": "COMPLETED"}), "R012"),
        ("execution-side-effect", lambda p: p.update({"real_side_effect_count": 1}), "R012"),
        ("execution-email-recipient", lambda p: p["company_email"].update({"recipient": "leaders@company.example.com"}), "R012"),
        ("execution-email-count", lambda p: p["company_email"].update({"eligible_count": 4}), "R012"),
        ("execution-email-drafts", lambda p: p["company_email"].update({"draft_count": 2}), "R012"),
        ("execution-email-sent", lambda p: p["company_email"].update({"sent_count": 1, "state": "SENT"}), "R012"),
        ("execution-manager-required", lambda p: p["lateral_manager_communications"].update({"required_count": 1}), "R010"),
        ("execution-manager-identity", lambda p: p["lateral_manager_communications"].update({"manager_identity_available_count": 2}), "R010"),
        ("execution-manager-draft", lambda p: p["lateral_manager_communications"].update({"draft_count": 2}), "R010"),
        ("execution-manager-sent", lambda p: p["lateral_manager_communications"].update({"sent_count": 2}), "R010"),
        ("execution-manager-state", lambda p: p["lateral_manager_communications"].update({"state": "SENT"}), "R010"),
        ("execution-pending-name", lambda p: p["pending_hold"].update({"employee": "Greta Nilsson"}), "R010"),
        ("execution-pending-draft", lambda p: p["pending_hold"].update({"announcement_draft_count": 1}), "R010"),
        ("execution-slack-draft", lambda p: p["slack"].update({"draft_count": 1}), "R012"),
        ("execution-slack-post", lambda p: p["slack"].update({"posted_count": 1, "state": "POSTED"}), "R012"),
        ("execution-wrong-channel", lambda p: p["slack"].update({"channel": "C_OTHER / private"}), "R012"),
        ("execution-remove-trace", lambda p: p.update({"source_trace": []}), "R015"),
    ]
    for name, mutator, criterion in execution_cases:
        cases.append((name, lambda root, m=mutator: mutate_json(root / EXECUTION, m), criterion))

    cases.extend([
        ("word-ava-title", lambda root: replace_docx_text(root / DOCX, "Senior Software Engineer", "Staff Software Engineer", row_token="Ava Chen"), "R005"),
        ("word-brandon-status", lambda root: replace_docx_text(root / DOCX, "ELIGIBLE_COMPANY_ANNOUNCEMENT", "PENDING_HOLD_NOT_CLEARED", row_token="Brandon Osei"), "R006"),
        ("word-elena-department", lambda root: replace_docx_text(root / DOCX, "Engineering", "Product", row_token="Elena Vasquez"), "R007"),
        ("word-greta-omitted", lambda root: replace_docx_text(root / DOCX, "Greta Nilsson", "Greta Missing", row_token="Greta Nilsson"), "R008"),
        ("word-hugo-row", lambda root: replace_docx_text(root / DOCX, "9", "10", row_token="Hugo Fernandez"), "R009"),
        ("word-clara-route", lambda root: replace_docx_text(root / DOCX, "RECEIVING_MANAGER_COMMUNICATION_REQUIRED", "ELIGIBLE_COMPANY_ANNOUNCEMENT", row_token="Clara Dubois"), "R010"),
        ("word-derek-hold", lambda root: replace_docx_text(root / DOCX, "YES", "NO", row_token="Derek Holm"), "R010"),
        ("word-faisal-manager", lambda root: replace_docx_text(root / DOCX, "UNAVAILABLE_IN_INPUT", "manager@company.example.com", row_token="Faisal Rahman"), "R010"),
        ("word-policy-lateral", lambda root: replace_docx_text(root / DOCX, "Lateral moves are NOT included", "Lateral moves are included", row_token="Lateral moves"), "R014"),
        ("word-emails-sent", lambda root: replace_docx_text(root / DOCX, "company emails sent=0", "company emails sent=1"), "R014"),
        ("word-manager-drafts", lambda root: replace_docx_text(root / DOCX, "manager drafts=0", "manager drafts=2"), "R014"),
        ("word-slack-posts", lambda root: replace_docx_text(root / DOCX, "Slack posts=0", "Slack posts=1"), "R014"),
        ("word-real-side-effects", lambda root: replace_docx_text(root / DOCX, "real side effects=0", "real side effects=1"), "R014"),
    ])

    cases.extend([
        ("ppt-drop-slide", lambda root: drop_ppt_slide(root / PPTX), "R013"),
        ("ppt-summary-count", lambda root: replace_pptx_text(root / PPTX, "five company eligible", "four company eligible"), "R013"),
        ("ppt-ava-name", lambda root: replace_pptx_text(root / PPTX, "Ava Chen", "Ava Wrong"), "R013"),
        ("ppt-brandon-name", lambda root: replace_pptx_text(root / PPTX, "Brandon Osei", "Brandon Wrong"), "R013"),
        ("ppt-elena-name", lambda root: replace_pptx_text(root / PPTX, "Elena Vasquez", "Elena Wrong"), "R013"),
        ("ppt-greta-name", lambda root: replace_pptx_text(root / PPTX, "Greta Nilsson", "Greta Missing"), "R013"),
        ("ppt-hugo-name", lambda root: replace_pptx_text(root / PPTX, "Hugo Fernandez", "Hugo Wrong"), "R013"),
        ("ppt-email-sent", lambda root: replace_pptx_text(root / PPTX, "email sent=0", "email sent=1"), "R013"),
        ("ppt-slack-posted", lambda root: replace_pptx_text(root / PPTX, "Slack drafts/posts=0", "Slack drafts/posts=1"), "R013"),
        ("ppt-greta-policy", lambda root: replace_pptx_text(root / PPTX, "Greta is included by full policy evaluation", "Greta is omitted"), "R008"),
        ("ppt-manager-invented", lambda root: replace_pptx_text(root / PPTX, "no manager identity is invented", "manager identity invented"), "R010"),
        ("ppt-clara-omitted", lambda root: replace_pptx_text(root / PPTX, "Clara Dubois", "Clara Missing"), "R013"),
        ("ppt-pending-cleared", lambda root: replace_pptx_text(root / PPTX, "PENDING_HOLD_NOT_CLEARED", "ELIGIBLE_COMPANY_ANNOUNCEMENT"), "R013"),
        ("ppt-authorized-slack", lambda root: replace_pptx_text(root / PPTX, "NOT_AUTHORIZED_CHANNEL", "POSTED"), "R013"),
        ("ppt-manager-drafts", lambda root: replace_pptx_text(root / PPTX, "0 identities · 0 drafts", "2 identities · 2 drafts"), "R013"),
        ("ppt-real-side-effects", lambda root: replace_pptx_text(root / PPTX, "Real side effects=0", "Real side effects=1"), "R013"),
    ])
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--json-out", type=Path)
    args = parser.parse_args()
    judge = load_judge()
    cases = build_cases()
    results: list[dict] = []
    failures: list[str] = []
    for name, mutator, expected_failure in cases:
        with tempfile.TemporaryDirectory(prefix=f"negative-5112-{name}-") as temporary:
            artifact_root = Path(temporary) / "artifacts"
            shutil.copytree(GOLDEN, artifact_root)
            try:
                mutator(artifact_root)
                result = judge.evaluate(artifact_root)
                detected = not result.get("criteria", {}).get(expected_failure, {}).get("passed", True)
                rejected = not result.get("passed", False) and float(result.get("reward", 1.0)) < 0.8
                results.append({"case": name, "expected_failure": expected_failure, "detected": detected, "rejected": rejected, "reward": result.get("reward"), "hurdle_failures": result.get("hurdle_failures", [])})
                if not detected or not rejected:
                    failures.append(f"{name}: expected={expected_failure}, detected={detected}, reward={result.get('reward')}")
            except Exception as exc:
                failures.append(f"{name}: mutation/evaluation error: {type(exc).__name__}: {exc}")
    payload = {
        "task_id": "cowork-evo/automationbench-5112-promotion-announcement-governance-package",
        "case_count": len(results), "declared_case_count": len(cases),
        "all_rejected": not failures and len(results) == len(cases),
        "max_reward": max((float(item["reward"]) for item in results), default=None),
        "failures": failures, "results": results,
    }
    if args.json_out:
        args.json_out.parent.mkdir(parents=True, exist_ok=True)
        args.json_out.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    if failures or len(results) != len(cases):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
