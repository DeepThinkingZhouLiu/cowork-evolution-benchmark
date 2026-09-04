#!/usr/bin/env python3
"""AutomationBench 5113 I-9 合规治理题的逐题确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import mailbox
import re
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from email.utils import getaddresses
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


TASK_ID = "cowork-evo/automationbench-5113-i9-compliance-control-package"
XLSX = "I9_Compliance_Audit.xlsx"
DOCX = "I9_Compliance_Control_Brief.docx"
PPTX = "I9_Compliance_Review.pptx"
MBOX = "I9_Notification_Drafts.mbox"
UPDATES = "I9_Tracker_Update_Drafts.json"
SLACK = "I9_Slack_Control.json"
EXECUTION = "I9_Execution_Control.json"
REQUIRED = {XLSX, DOCX, PPTX, MBOX, UPDATES, SLACK, EXECUTION}
INPUT_HASHES = {
    "I9_Verification_Tracker.xlsx": "89133a81d4733ee8945f4d735cfd8261b2b1fb12eadf004739e877e278d8e72c",
    "Employment_Verification_Policy.mbox": "c6c27b85e2eb98e4fd4c61183ea038f0dd70c647e979ed18e0cbf81fad62e848",
    "Slack_Channels.json": "d4d3895e0acb156a5d271b03316f63e06b9bc4f27997f6b6e41ba47b463025b2",
    "Workflow_Context.json": "76c4a0657bc27f2fd830b10613ec7ceae3adcdc5798aeee0f822d27bf69bbc99",
}
WEIGHTS = {"R001": 4, "R002": 6, "R003": 6, "R004": 7, "R005": 4, "R006": 4, "R007": 4, "R008": 5, "R009": 5, "R010": 5, "R011": 5, "R012": 5, "R013": 9, "R014": 8, "R015": 7, "R016": 6, "R017": 4, "R018": 6}
HURDLE_CAPS = {"R001": 0.0, "R002": 0.2, "R003": 0.2, "R004": 0.3, "R005": 0.3, "R006": 0.3, "R007": 0.3, "R008": 0.3, "R009": 0.3, "R010": 0.3, "R011": 0.25, "R012": 0.25, "R013": 0.2, "R014": 0.2, "R015": 0.3, "R016": 0.35}
SOURCE = [
    (2, "Amara Obi", "amara.obi@company.example.com", "mgr-eng@company.example.com", "2026-04-14", "Incomplete", ""),
    (3, "Boris Petrov", "boris.petrov@company.example.com", "mgr-sales@company.example.com", "2026-04-07", "Incomplete", ""),
    (4, "Cassie Lin", "cassie.lin@company.example.com", "mgr-eng@company.example.com", "2026-04-10", "Complete", ""),
    (5, "Dmitri Volkov", "dmitri.volkov@company.example.com", "mgr-mkt@company.example.com", "2026-04-01", "Incomplete", ""),
    (6, "Esme Nakamura", "esme.nakamura@company.example.com", "mgr-sales@company.example.com", "2026-04-11", "Incomplete", ""),
    (7, "Frankie Reeves", "frankie.reeves@company.example.com", "mgr-eng@company.example.com", "2026-03-25", "Incomplete", ""),
    (8, "Gina Torres", "gina.torres@company.example.com", "mgr-mkt@company.example.com", "2026-04-15", "Incomplete", ""),
    (9, "Hank Morrison", "hank.morrison@company.example.com", "mgr-sales@company.example.com", "2026-04-08", "Incomplete", ""),
]
TRUTH = {
    "Amara Obi": ("2026-04-17", 1, "NO", "YES", "YES", "NO", "NO", "Reminder Sent", "YES", "NO"),
    "Boris Petrov": ("2026-04-10", 6, "NO", "NO", "NO", "YES", "NO", "Violation - Legal Notified", "YES", "NO"),
    "Cassie Lin": ("2026-04-15", 3, "YES", "NO", "NO", "NO", "YES", "Complete", "NO", "NO"),
    "Dmitri Volkov": ("2026-04-06", 10, "NO", "NO", "NO", "YES", "NO", "Violation - Legal Notified", "YES", "NO"),
    "Esme Nakamura": ("2026-04-15", 3, "NO", "YES", "YES", "NO", "NO", "Reminder Sent", "YES", "NO"),
    "Frankie Reeves": ("2026-03-30", 15, "NO", "NO", "NO", "YES", "NO", "Violation - Legal Notified", "YES", "NO"),
    "Gina Torres": ("2026-04-20", 0, "NO", "YES", "YES", "NO", "NO", "Reminder Sent", "YES", "NO"),
    "Hank Morrison": ("2026-04-13", 5, "NO", "NO", "NO", "YES", "NO", "Violation - Legal Notified", "YES", "NO"),
}
CRITERION = {"Amara Obi": "R005", "Esme Nakamura": "R006", "Gina Torres": "R007", "Boris Petrov": "R008", "Dmitri Volkov": "R009", "Frankie Reeves": "R010", "Hank Morrison": "R011", "Cassie Lin": "R012"}
POLICIES = [
    (1, "Overview", "All new hires must complete Form I-9 verification."),
    (2, "Employee documents", "List A establishes identity and authorization; List B + List C establish them separately."),
    (3, "Employer deadline", "Complete Section 2 within 3 business days of the employee's start date."),
    (4, "Incomplete within window", "Send employee an 'I-9 Reminder', CC the manager, and update Status to 'Reminder Sent'."),
    (5, "Incomplete past window", "Email legal@company.example.com with employee name and start date; update Status to 'Violation - Legal Notified'."),
    (6, "Record retention", "Retain I-9 forms for 3 years after hire or 1 year after termination, whichever is later."),
    (7, "Audits", "Internal audits are conducted quarterly."),
]
REMINDERS = [SOURCE[0], SOURCE[4], SOURCE[6]]
VIOLATIONS = [SOURCE[1], SOURCE[3], SOURCE[5], SOURCE[7]]


@dataclass
class Ledger:
    values: dict[str, dict[str, Any]]

    def __init__(self) -> None:
        self.values = {key: {"passed": True, "details": []} for key in WEIGHTS}

    def require(self, key: str, condition: bool, detail: str) -> None:
        if not condition:
            self.values[key]["passed"] = False
            self.values[key]["details"].append(detail)

    def fail(self, key: str, detail: str) -> None:
        self.require(key, False, detail)

    def passed(self, key: str) -> bool:
        return bool(self.values[key]["passed"])


def root() -> Path:
    return Path(__file__).resolve().parents[1]


def input_root() -> Path:
    container = Path("/app/input_files")
    return container if container.is_dir() else root() / "data/input_files"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def norm(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def rows(sheet, start: int, end: int, columns: int) -> list[tuple[Any, ...]]:
    return [tuple(norm(sheet.cell(row, column).value) for column in range(1, columns + 1)) for row in range(start, end + 1)]


def formula(value: Any, *tokens: str) -> bool:
    text = str(value or "")
    return text.startswith("=") and all(token.casefold() in text.casefold() for token in tokens)


def decode(message) -> str:
    if message.is_multipart():
        parts = []
        for part in message.walk():
            if part.get_content_type() == "text/plain":
                parts.append((part.get_payload(decode=True) or b"").decode(part.get_content_charset() or "utf-8", errors="replace"))
        return "\n".join(parts).strip()
    return (message.get_payload(decode=True) or b"").decode(message.get_content_charset() or "utf-8", errors="replace").strip()


def addresses(message, header: str) -> set[str]:
    return {address.casefold() for _, address in getaddresses(message.get_all(header, [])) if address}


def source_policy() -> dict[str, str]:
    box = mailbox.mbox(input_root() / "Employment_Verification_Policy.mbox", create=False)
    messages = list(box)
    box.close()
    if len(messages) != 1:
        raise ValueError("源政策邮件数量错误")
    message = messages[0]
    return {"message_id": str(message.get("X-AutomationBench-Message-ID", "")), "from": str(message.get("From", "")), "to": str(message.get("To", "")), "subject": str(message.get("Subject", "")), "date": str(message.get("Date", "")), "body": decode(message)}


def trace(source_row: int, *, frozen: bool = True) -> str:
    value = f"I9_Verification_Tracker.xlsx:Recent Hires row {source_row} | Employment_Verification_Policy.mbox:msg_i9_policy_5113"
    return value + " | Workflow_Context.json:meta.current_time" if frozen else value


def expected_decisions() -> dict[str, tuple[Any, ...]]:
    result = {}
    for source_row, employee, employee_email, manager_email, start_date, status, notes in SOURCE:
        # LibreOffice 对“公式引用空白单元格”的缓存值为数值 0；源表本身仍严格保留空白字符串。
        result[employee] = (source_row, employee, employee_email, manager_email, start_date, status, 0 if notes == "" else notes, *TRUTH[employee], trace(source_row))
    return result


def render(path: Path) -> tuple[int, str]:
    with tempfile.TemporaryDirectory(prefix="judge-5113-") as temporary:
        work = Path(temporary)
        output = work / "out"
        output.mkdir()
        process = subprocess.run(["libreoffice", "--headless", f"-env:UserInstallation=file://{work / 'profile'}", "--convert-to", "pdf", "--outdir", str(output), str(path)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300)
        pdf = output / f"{path.stem}.pdf"
        if process.returncode != 0 or not pdf.is_file():
            raise RuntimeError(process.stdout[-1000:])
        info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True, timeout=60)
        match = re.search(r"^Pages:\s+(\d+)", info.stdout, flags=re.M)
        text = subprocess.run(["pdftotext", str(pdf), "-"], capture_output=True, text=True, timeout=60)
        if info.returncode != 0 or not match or text.returncode != 0:
            raise RuntimeError("无法读取渲染 PDF")
        return int(match.group(1)), text.stdout


def workbook_checks(artifacts: Path, checks: Ledger) -> dict[str, tuple[Any, ...]]:
    formulas = load_workbook(artifacts / XLSX, data_only=False)
    values = load_workbook(artifacts / XLSX, data_only=True)
    names = ["Source_Hires", "Policy_Source", "Policy_Rules", "Evidence_Register", "Compliance_Decisions", "Reminder_Queue", "Legal_Violation_Queue", "No_Action_Queue", "Control_Summary"]
    checks.require("R001", formulas.sheetnames == names, "Excel 九张工作表名称或顺序错误")
    checks.require("R002", rows(values["Source_Hires"], 2, 9, 7) == SOURCE, "Source_Hires 未逐行保留八条上游记录")

    policy = source_policy()
    expected_policy = (policy["message_id"], policy["from"], policy["to"], policy["subject"], policy["date"], policy["body"], "Employment_Verification_Policy.mbox:msg_i9_policy_5113")
    checks.require("R003", rows(values["Policy_Source"], 2, 2, 7) == [expected_policy], "Legal 政策邮件全文或元数据错误")
    expected_rules = [(order, rule, requirement, "Employment_Verification_Policy.mbox:msg_i9_policy_5113") for order, rule, requirement in POLICIES]
    checks.require("R003", rows(values["Policy_Rules"], 2, 8, 4) == expected_rules, "七条政策规则错误")
    evidence = rows(values["Evidence_Register"], 2, 12, 5)
    checks.require("R003", len(evidence) == 11, "证据表必须恰好 11 条")
    for source, actual in zip(SOURCE, evidence[:8]):
        source_row, employee, employee_email, manager_email, start_date, status, notes = source
        expected_value = f"{employee} | {employee_email} | {manager_email} | {start_date} | {status} | {notes}"
        checks.require("R003", actual == (f"E-ROW-{source_row}", "Source tracker row", employee, expected_value, f"I9_Verification_Tracker.xlsx:Recent Hires row {source_row}"), f"{employee} 证据错误")
    checks.require("R003", evidence[8] == ("E-msg_i9_policy_5113", "Authoritative email", f"{policy['from']} | {policy['subject']}", policy["body"], "Employment_Verification_Policy.mbox:msg_i9_policy_5113"), "权威政策证据错误")
    checks.require("R003", evidence[9] == ("E-TIME", "Frozen runtime", "Workflow context", "2026-04-15T09:00:00Z", "Workflow_Context.json:meta.current_time"), "冻结时间证据错误")
    checks.require("R003", evidence[10] == ("E-CHANNEL-C_HR_5113", "Channel snapshot", "C_HR_5113 / hr-general", '{"id":"C_HR_5113","name":"hr-general","is_private":false}', "Slack_Channels.json:C_HR_5113"), "Slack 频道证据错误")

    decision_formulas = formulas["Compliance_Decisions"]
    for row_index in range(2, 10):
        for column in (1, 2, 3, 4, 6, 7):
            letter = chr(64 + column)
            checks.require("R004", formula(decision_formulas.cell(row_index, column).value, f"Source_Hires!{letter}{row_index}"), f"Compliance_Decisions!{letter}{row_index} 未链接来源")
        dependencies = {5: (f"Source_Hires!E{row_index}", "DATEVALUE"), 8: (f"E{row_index}", "WORKDAY", "3"), 9: (f"E{row_index}+1", "NETWORKDAYS", "DATE(2026,4,15)"), 10: (f"F{row_index}", "Complete"), 11: (f"J{row_index}", f"I{row_index}", "<=3"), 12: (f"K{row_index}",), 13: (f"J{row_index}", f"K{row_index}"), 14: (f"J{row_index}",), 15: (f"L{row_index}", f"M{row_index}"), 16: (f"L{row_index}", f"M{row_index}"), 17: ("NO",), 18: (f"A{row_index}", "msg_i9_policy_5113", "current_time")}
        for column, tokens in dependencies.items():
            checks.require("R004", formula(decision_formulas.cell(row_index, column).value, *tokens), f"Compliance_Decisions 第 {row_index} 行第 {column} 列公式依赖错误")
    for sheet_name, end_row, end_col in (("Reminder_Queue", 4, 12), ("Legal_Violation_Queue", 5, 11), ("No_Action_Queue", 2, 8)):
        sheet = formulas[sheet_name]
        for row_index in range(2, end_row + 1):
            for column in range(2, end_col + 1):
                checks.require("R004", str(sheet.cell(row_index, column).value or "").startswith("="), f"{sheet_name}!{sheet.cell(row_index, column).coordinate} 被静态化")
    for row_index in range(2, 15):
        checks.require("R004", str(formulas["Control_Summary"].cell(row_index, 2).value or "").startswith("="), f"Control_Summary!B{row_index} 被静态化")

    actual = {str(item[1]): item for item in rows(values["Compliance_Decisions"], 2, 9, 18)}
    wanted = expected_decisions()
    checks.require("R002", list(actual) == [item[1] for item in SOURCE], "决策表人员集合或来源顺序错误")
    for employee, expected in wanted.items():
        checks.require(CRITERION[employee], actual.get(employee) == expected, f"{employee} 的日期、工作日、状态或路由错误")

    reminder_expected = []
    for order, source in enumerate(REMINDERS, start=1):
        source_row, employee, employee_email, manager_email, start_date, _, _ = source
        deadline, elapsed = TRUTH[employee][:2]
        reminder_expected.append((order, source_row, employee, employee_email, manager_email, start_date, deadline, elapsed, "I-9 Reminder", "DRAFT_NOT_SENT", False, trace(source_row, frozen=False)))
    checks.require("R013", rows(values["Reminder_Queue"], 2, 4, 12) == reminder_expected, "Excel 提醒队列的三人、经理 Cc、日期、顺序或状态错误")
    legal_expected = []
    for order, source in enumerate(VIOLATIONS, start=1):
        source_row, employee, _, _, start_date, _, _ = source
        deadline, elapsed = TRUTH[employee][:2]
        legal_expected.append((order, source_row, employee, start_date, deadline, elapsed, "legal@company.example.com", f"I-9 Compliance Violation - {employee}", "DRAFT_NOT_SENT", False, trace(source_row, frozen=False)))
    checks.require("R013", rows(values["Legal_Violation_Queue"], 2, 5, 11) == legal_expected, "Excel 法务违规队列的四人、Hank、日期、顺序或状态错误")
    no_action_expected = [(1, 4, "Cassie Lin", "Complete", "SOURCE_STATUS_COMPLETE", False, False, trace(4, frozen=False))]
    checks.require("R012", rows(values["No_Action_Queue"], 2, 2, 8) == no_action_expected, "Cassie 无动作队列错误")
    checks.require("R004", [norm(values["Control_Summary"].cell(row_index, 2).value) for row_index in range(2, 15)] == [8, 7, 1, 3, 4, 1, 7, 0, 7, 0, 0, 0, 0], "Excel 十三项汇总缓存值错误")
    total_tables = sum(len(sheet.tables) for sheet in formulas.worksheets)
    cf_rules = sum(len(sheet.conditional_formatting[item]) for sheet in formulas.worksheets for item in sheet.conditional_formatting)
    charts = sum(len(sheet._charts) for sheet in formulas.worksheets)
    checks.require("R018", total_tables == 9, f"Excel 应有 9 张原生表，实际 {total_tables}")
    checks.require("R018", cf_rules >= 3, "Excel 缺少三类决策条件格式")
    checks.require("R018", charts >= 1, "Excel 缺少可编辑原生图表")
    with zipfile.ZipFile(artifacts / XLSX) as archive:
        names_in_zip = archive.namelist()
        checks.require("R001", not any(name.endswith("vbaProject.bin") for name in names_in_zip), "Excel 包含宏")
        checks.require("R001", not any(name.startswith("xl/externalLinks/") for name in names_in_zip), "Excel 包含外部链接")
    formulas.close()
    values.close()
    return actual


def email_body(source: tuple[Any, ...]) -> str:
    source_row, employee, employee_email, manager_email, start_date, _, _ = source
    deadline, elapsed, _, _, reminder, _, _, _, _, _ = TRUTH[employee]
    if reminder == "YES":
        return "\n".join([f"Hello {employee},", "", f"Our records show that your Form I-9 verification is incomplete. Start Date={start_date}; third business-day deadline={deadline}; elapsed business days={elapsed}.", "Please complete the required verification within the policy window.", "", "DRAFT_NOT_SENT — this offline benchmark artifact has not been sent.", "- HR Operations"])
    return "\n".join(["Legal Team,", "", f"I-9 compliance violation: Employee={employee}; Start Date={start_date}; third business-day deadline={deadline}; elapsed business days={elapsed}.", "The source tracker remains Incomplete and the 3-business-day window has passed.", "", "DRAFT_NOT_SENT — this offline benchmark artifact has not been sent.", "- HR Operations"])


def mbox_checks(artifacts: Path, checks: Ledger) -> list[str]:
    box = mailbox.mbox(artifacts / MBOX, create=False)
    messages = list(box)
    box.close()
    actions = [item for item in SOURCE if item[1] != "Cassie Lin"]
    checks.require("R013", len(messages) == 7, "MBOX 必须恰好包含 7 封草稿")
    names = []
    if len(messages) != 7:
        return names
    checks.require("R013", [str(message.get("Message-ID", "")) for message in messages] == [f"<i9-5113-{item[0]}@evocowork.local>" for item in actions], "MBOX 七封草稿未保持来源行顺序")
    for order, (message, source) in enumerate(zip(messages, actions), start=1):
        source_row, employee, employee_email, manager_email, _, _, _ = source
        reminder = employee in {item[1] for item in REMINDERS}
        criterion = CRITERION[employee]
        names.append(employee)
        checks.require(criterion, str(message.get("From", "")).casefold() == "hr-ops@company.example.com", f"{employee} 邮件发件人错误")
        checks.require(criterion, addresses(message, "To") == ({employee_email} if reminder else {"legal@company.example.com"}), f"{employee} 邮件收件人错误")
        checks.require(criterion, addresses(message, "Cc") == ({manager_email} if reminder else set()) and not addresses(message, "Bcc"), f"{employee} 邮件 Cc/Bcc 错误")
        expected_subject = "I-9 Reminder" if reminder else f"I-9 Compliance Violation - {employee}"
        checks.require(criterion, str(message.get("Subject", "")) == expected_subject, f"{employee} 邮件主题错误")
        checks.require(criterion, str(message.get("Date", "")) == "Wed, 15 Apr 2026 09:00:00 +0000", f"{employee} 邮件冻结日期错误")
        checks.require(criterion, str(message.get("Message-ID", "")) == f"<i9-5113-{source_row}@evocowork.local>" and str(message.get("X-EvoCowork-Source-Row", "")) == str(source_row), f"{employee} 邮件来源行或 Message-ID 错误")
        expected_route = "REMINDER_WITH_MANAGER_CC" if reminder else "VIOLATION_LEGAL_NOTIFICATION"
        checks.require(criterion, str(message.get("X-EvoCowork-Status", "")) == "DRAFT_NOT_SENT" and str(message.get("X-EvoCowork-Route", "")) == expected_route, f"{employee} 邮件草稿状态或路由错误")
        checks.require(criterion, decode(message) == email_body(source), f"{employee} 邮件正文事实或零发送声明错误")
        checks.require("R013", not message.is_multipart() or all(part.get_content_disposition() != "attachment" for part in message.walk()), f"{employee} 邮件不应包含附件")
    checks.require("R012", "Cassie Lin" not in "\n".join(decode(message) for message in messages), "Cassie 错误出现在邮件草稿中")
    return names


def expected_update(source: tuple[Any, ...]) -> dict[str, Any]:
    source_row, employee, _, _, start_date, status, _ = source
    deadline, elapsed, _, _, reminder, _, _, proposed, _, _ = TRUTH[employee]
    return {"source_row": source_row, "employee": employee, "start_date": start_date, "deadline": deadline, "elapsed_business_days": elapsed, "source_status": status, "route": "REMINDER_WITH_MANAGER_CC" if reminder == "YES" else "VIOLATION_LEGAL_NOTIFICATION", "proposed_status": proposed, "state": "DRAFT_NOT_APPLIED", "applied": False, "source_trace": [f"I9_Verification_Tracker.xlsx:Recent Hires row {source_row}", "Employment_Verification_Policy.mbox:msg_i9_policy_5113", "Workflow_Context.json:meta.current_time"]}


def update_checks(artifacts: Path, checks: Ledger) -> list[dict[str, Any]]:
    payload = json.loads((artifacts / UPDATES).read_text(encoding="utf-8"))
    expected_drafts = [expected_update(item) for item in SOURCE if item[1] != "Cassie Lin"]
    expected = {"task_id": "automationbench-5113", "status": "CONTROLLED_DRAFT_ONLY", "frozen_time": "2026-04-15T09:00:00Z", "business_day_method": "Monday-Friday; count strictly after Start Date through frozen date; no unstated holidays", "draft_count": 7, "applied_count": 0, "drafts": expected_drafts}
    checks.require("R013", payload == expected, "Tracker 更新草稿的人员、顺序、日期、路由、状态或未应用边界错误")
    checks.require("R012", all(item.get("employee") != "Cassie Lin" for item in payload.get("drafts", []) if isinstance(item, dict)), "Cassie 错误出现在更新草稿中")
    return payload.get("drafts", []) if isinstance(payload.get("drafts"), list) else []


def slack_checks(artifacts: Path, checks: Ledger) -> dict[str, Any]:
    payload = json.loads((artifacts / SLACK).read_text(encoding="utf-8"))
    expected = {"task_id": "automationbench-5113", "source_channel": {"id": "C_HR_5113", "name": "hr-general", "is_private": False}, "decision": "NOT_AUTHORIZED_CHANNEL", "reason": "The Legal employment verification guide authorizes employee email with manager Cc or legal-team email; it does not authorize Slack for I-9 status.", "draft_count": 0, "posted_count": 0, "posted": False, "employee_names_in_slack": [], "i9_statuses_in_slack": [], "messages_created": [], "source_trace": ["Employment_Verification_Policy.mbox:msg_i9_policy_5113", "Slack_Channels.json:C_HR_5113", "I9_Compliance_Audit.xlsx:Control_Summary"]}
    checks.require("R014", payload == expected, "Slack 未保持未授权、零草稿、零发布和零 I-9 信息泄露")
    return payload


def execution_checks(artifacts: Path, checks: Ledger) -> dict[str, Any]:
    payload = json.loads((artifacts / EXECUTION).read_text(encoding="utf-8"))
    expected = {"task_id": "automationbench-5113", "status": "CONTROLLED_DRAFT_ONLY", "email": {"reminder_drafts": 3, "legal_violation_drafts": 4, "total_drafts": 7, "sent_count": 0, "state": "DRAFT_NOT_SENT"}, "tracker_updates": {"reminder_status_drafts": 3, "violation_status_drafts": 4, "total_drafts": 7, "applied_count": 0, "state": "DRAFT_NOT_APPLIED"}, "no_action": {"employee": "Cassie Lin", "email_draft_count": 0, "tracker_update_draft_count": 0, "state": "SOURCE_STATUS_COMPLETE"}, "slack": {"channel": "C_HR_5113 / hr-general", "draft_count": 0, "posted_count": 0, "state": "NOT_AUTHORIZED_CHANNEL"}, "real_side_effect_count": 0, "source_trace": ["I9_Verification_Tracker.xlsx:Recent Hires rows 2-9", "Employment_Verification_Policy.mbox:msg_i9_policy_5113", "Workflow_Context.json:meta.current_time", "Slack_Channels.json:C_HR_5113"]}
    checks.require("R014", payload == expected, "执行控制的邮件、台账、Slack 或真实副作用边界错误")
    return payload


def word_checks(artifacts: Path, checks: Ledger) -> str:
    document = Document(artifacts / DOCX)
    cell_rows = [tuple(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
    text = "\n".join(paragraph.text for paragraph in document.paragraphs) + "\n" + "\n".join(" | ".join(row) for row in cell_rows)
    checks.require("R016", len(document.tables) == 5, f"Word 应有 5 张原生表，实际 {len(document.tables)}")
    for values in [("Reminder", "3", "Amara; Esme; Gina", "Employee + manager Cc", "3 drafts / 0 sent"), ("Violation", "4", "Boris; Dmitri; Frankie; Hank", "legal@company.example.com", "4 drafts / 0 sent"), ("No action", "1", "Cassie", "Source status Complete", "0 email / 0 update"), ("Slack", "0", "No employee or status", "NOT_AUTHORIZED_CHANNEL", "0 draft / 0 post")]:
        checks.require("R016", values in cell_rows, f"Word 汇总行 {values[0]} 错误")
    for order, rule, requirement in POLICIES:
        checks.require("R016", (str(order), rule, requirement) in cell_rows, f"Word 缺少政策 {order}")
    for source in SOURCE:
        source_row, employee, employee_email, manager_email, start_date, status, _ = source
        deadline, elapsed, _, within, reminder, legal, _, proposed, _, _ = TRUTH[employee]
        expected = (str(source_row), employee, employee_email, manager_email, start_date, deadline, str(elapsed), status, within, reminder, legal, proposed)
        checks.require(CRITERION[employee], expected in cell_rows, f"Word 中 {employee} 决策行错误")
    for source in [item for item in SOURCE if item[1] != "Cassie Lin"]:
        _, employee, employee_email, manager_email, start_date, _, _ = source
        reminder = employee in {item[1] for item in REMINDERS}
        expected = (employee, "REMINDER" if reminder else "LEGAL", employee_email if reminder else "legal@company.example.com", manager_email if reminder else "", start_date, TRUTH[employee][7], "1 / 0")
        checks.require("R016", expected in cell_rows, f"Word 中 {employee} 草稿队列错误")
    for values in [("Emails", "7", "0", "DRAFT_NOT_SENT"), ("Tracker updates", "7", "0", "DRAFT_NOT_APPLIED"), ("Slack", "0", "0", "NOT_AUTHORIZED_CHANNEL"), ("Real side effects", "0", "0", "OFFLINE_ONLY")]:
        checks.require("R016", values in cell_rows, f"Word 执行控制 {values[0]} 错误")
    required = ["three REMINDER_WITH_MANAGER_CC", "four VIOLATION_LEGAL_NOTIFICATION", "one NO_ACTION_COMPLETE", "Hank Morrison", "upstream assertions do not name him", "Cassie Lin is Complete", "Monday through Friday", "counted strictly after Start Date", "Amara=1", "Boris=6", "Cassie=3", "Dmitri=10", "Esme=3", "Frankie=15", "Gina=0", "Hank=5", "source hires=8", "incomplete=7", "complete=1", "reminders=3", "violations=4", "no action=1", "real side effects=0"]
    checks.require("R016", all(token in text for token in required), "Word 缺少工作日方法、Hank 补全、Cassie 保护、计数或零执行证据")
    checks.require("R016", not re.search(r"(?:emails sent|tracker updates applied|Slack posts|real side effects)\s*=\s*[1-9]", text, flags=re.I), "Word 出现与离线草稿边界冲突的已执行声明")
    with zipfile.ZipFile(artifacts / DOCX) as archive:
        names_in_zip = archive.namelist()
        checks.require("R001", not any(name.endswith("vbaProject.bin") for name in names_in_zip), "Word 包含宏")
        checks.require("R001", not any(name.startswith("word/embeddings/") for name in names_in_zip), "Word 包含嵌入对象")
    return text


def ppt_checks(artifacts: Path, checks: Ledger) -> str:
    presentation = Presentation(artifacts / PPTX)
    checks.require("R015", len(presentation.slides) == 3, "PPT 必须恰好 3 页")
    texts = []
    table_counts = []
    pictures = 0
    for slide in presentation.slides:
        parts = []
        table_count = 0
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                parts.append(shape.text.strip())
            if getattr(shape, "has_table", False):
                table_count += 1
                parts.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
        texts.append("\n".join(parts))
        table_counts.append(table_count)
    checks.require("R015", pictures == 0, "PPT 关键内容不得以图片替代")
    checks.require("R015", table_counts == [1, 1, 1], "PPT 每页应恰有一张原生表")
    if len(texts) >= 1:
        checks.require("R015", all(token in texts[0] for token in ("Eight reviewed", "three reminders", "four violations", "one no action", "zero side effects", "REMINDER | 3 | Amara; Esme; Gina", "LEGAL | 4 | Boris; Dmitri; Frankie; Hank", "NO ACTION | 1 | Cassie", "Email drafts=7", "emails sent=0", "updates applied=0", "Hank is included", "Cassie is protected")), "PPT 第 1 页汇总或执行边界错误")
    if len(texts) >= 2:
        for source in SOURCE:
            source_row, employee, _, _, start_date, status, _ = source
            deadline, elapsed, _, within, reminder, legal, _, proposed, _, _ = TRUTH[employee]
            expected = f"{source_row} | {employee} | {start_date} | {deadline} | {elapsed} | {status} | {within} | {reminder} | {legal} | {proposed}"
            checks.require(CRITERION[employee], expected in texts[1], f"PPT 第 2 页 {employee} 决策错误")
        checks.require("R015", all(token in texts[1] for token in ("Weekdays only", "count strictly after Start Date", "Esme remains within on day 3", "Gina starts today (0 elapsed)", "Hank is a 5-day violation")), "PPT 第 2 页工作日边界说明错误")
    if len(texts) >= 3:
        for order, source in enumerate([item for item in SOURCE if item[1] != "Cassie Lin"], start=1):
            _, employee, employee_email, manager_email, start_date, _, _ = source
            reminder = employee in {item[1] for item in REMINDERS}
            expected = f"{order} | {employee} | {'REMINDER' if reminder else 'LEGAL'} | {employee_email if reminder else 'legal@company.example.com'} | {manager_email if reminder else ''} | {start_date} | DRAFT_NOT_SENT"
            checks.require(CRITERION[employee], expected in texts[2], f"PPT 第 3 页 {employee} 草稿队列错误")
        checks.require("R015", all(token in texts[2] for token in ("Seven source-ordered drafts", "no email sent", "no tracker update applied", "no Slack message", "7 drafts · 0 applied", "NOT_AUTHORIZED_CHANNEL · 0 posts", "Real side effects=0")), "PPT 第 3 页队列或零执行控制错误")
    with zipfile.ZipFile(artifacts / PPTX) as archive:
        names_in_zip = archive.namelist()
        checks.require("R001", not any(name.endswith("vbaProject.bin") for name in names_in_zip), "PPT 包含宏")
        checks.require("R001", not any(name.startswith("ppt/embeddings/") for name in names_in_zip), "PPT 包含嵌入对象")
    return "\n".join(texts)


def evaluate(artifacts: Path) -> dict[str, Any]:
    checks = Ledger()
    for name, expected in INPUT_HASHES.items():
        path = input_root() / name
        checks.require("R001", path.is_file() and sha256(path) == expected, f"冻结输入 {name} 缺失或哈希错误")
    files = {path.name for path in artifacts.iterdir() if path.is_file()} if artifacts.is_dir() else set()
    checks.require("R001", files == REQUIRED, f"交付文件集合错误：{sorted(files)}")
    if not REQUIRED.issubset(files):
        for key in WEIGHTS:
            if key != "R001":
                checks.fail(key, "核心交付缺失")
        return finalize(checks, {}, {})
    for name in REQUIRED:
        checks.require("R001", (artifacts / name).stat().st_size > 0, f"{name} 为空")
    values: dict[str, Any] = {}
    calls = [
        ("workbook", workbook_checks, ("R001", "R002", "R003", "R004", "R005", "R006", "R007", "R008", "R009", "R010", "R011", "R012", "R013", "R017", "R018")),
        ("mbox", mbox_checks, ("R001", "R005", "R006", "R007", "R008", "R009", "R010", "R011", "R012", "R013", "R017")),
        ("updates", update_checks, ("R001", "R005", "R006", "R007", "R008", "R009", "R010", "R011", "R012", "R013", "R017")),
        ("slack", slack_checks, ("R001", "R014", "R017")),
        ("execution", execution_checks, ("R001", "R014", "R017")),
        ("word", word_checks, ("R001", "R005", "R006", "R007", "R008", "R009", "R010", "R011", "R012", "R016", "R017")),
        ("ppt", ppt_checks, ("R001", "R005", "R006", "R007", "R008", "R009", "R010", "R011", "R012", "R015", "R017")),
    ]
    for key, call, criteria in calls:
        try:
            values[key] = call(artifacts, checks)
        except Exception as exc:
            values[key] = {}
            for criterion in criteria:
                checks.fail(criterion, f"{key} 解析/核验失败：{exc}")
    if all(values.get(key) for key in ("workbook", "mbox", "updates", "slack", "execution", "word", "ppt")):
        expected_order = [item[1] for item in SOURCE if item[1] != "Cassie Lin"]
        checks.require("R017", values["mbox"] == expected_order, "MBOX 与来源动作顺序冲突")
        checks.require("R017", [item.get("employee") for item in values["updates"]] == expected_order, "Tracker JSON 与来源动作顺序冲突")
        checks.require("R017", all(name in values["word"] and name in values["ppt"] for _, name, *_ in SOURCE), "Word/PPT 八人覆盖冲突")
        checks.require("R017", values["execution"].get("real_side_effect_count") == 0 and values["slack"].get("posted_count") == 0, "Execution 与 Slack 零副作用冲突")
    pages = {}
    if checks.passed("R001"):
        for name, expected_pages in ((XLSX, 9), (DOCX, 3), (PPTX, 3)):
            try:
                page_count, rendered_text = render(artifacts / name)
                pages[name] = page_count
                checks.require("R018", page_count == expected_pages, f"{name} 应渲染 {expected_pages} 页，实际 {page_count}")
                checks.require("R018", not re.search(r"#(?:VALUE|N/A|REF|NAME|DIV/0)", rendered_text, flags=re.I), f"{name} 渲染出现公式错误")
            except Exception as exc:
                checks.fail("R018", f"{name} 渲染失败：{exc}")
    else:
        checks.fail("R018", "核心交付未通过")
    return finalize(checks, values.get("workbook", {}), pages)


def finalize(checks: Ledger, decisions: dict[str, Any], pages: dict[str, int]) -> dict[str, Any]:
    hurdles = [key for key in HURDLE_CAPS if not checks.passed(key)]
    base = sum(weight for key, weight in WEIGHTS.items() if checks.passed(key)) / 100
    cap = min([1.0] + [HURDLE_CAPS[key] for key in hurdles])
    penalties = []
    total_penalty = 0.0
    if any(not checks.passed(key) for key in ("R002", "R003", "R005", "R006", "R007", "R008", "R009", "R010", "R011")):
        penalties.append({"id": "R019", "amount": 0.3, "reason": "来源员工、邮箱、日期、状态、工作日、路由、顺序或汇总错误"})
        total_penalty += 0.3
    if any(not checks.passed(key) for key in ("R005", "R006", "R007", "R008", "R009", "R010", "R011", "R012", "R013")):
        penalties.append({"id": "R020", "amount": 0.3, "reason": "漏掉 Hank、误处理 Cassie、提醒未 Cc 经理或违规未通知法务"})
        total_penalty += 0.3
    if any(not checks.passed(key) for key in ("R013", "R014", "R015", "R016")):
        penalties.append({"id": "R021", "amount": 0.3, "reason": "Slack 泄露、邮件或台账虚假执行，或 Office 文件声称发生真实动作"})
        total_penalty += 0.3
    if any(not checks.passed(key) for key in ("R001", "R004", "R017", "R018")):
        penalties.append({"id": "R022", "amount": 0.25, "reason": "文件、公式、原生结构、追踪、跨文件或渲染错误"})
        total_penalty += 0.25
    reward = round(max(0.0, min(base, cap) - total_penalty), 6)
    return {"task_id": TASK_ID, "reward": reward, "passed": reward >= 0.8 and not hurdles, "base_score": round(base, 6), "hurdle_cap": cap, "hurdle_failures": hurdles, "penalties": penalties, "criteria": checks.values, "source_truth": {"input_sha256": INPUT_HASHES, "employees": 8, "incomplete": 7, "complete": 1, "reminders": 3, "violations": 4, "no_action": 1, "email_drafts": 7, "emails_sent": 0, "tracker_update_drafts": 7, "tracker_updates_applied": 0, "slack_posts": 0, "real_side_effects": 0}, "decision_count": len(decisions), "render_pages": pages}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    try:
        result = evaluate(args.artifact_dir)
    except Exception as exc:
        result = {"task_id": TASK_ID, "reward": 0.0, "passed": False, "fatal_error": f"{type(exc).__name__}: {exc}", "hurdle_failures": ["R001"]}
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{float(result['reward']):.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
