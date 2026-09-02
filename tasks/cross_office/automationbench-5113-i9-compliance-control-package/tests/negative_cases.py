#!/usr/bin/env python3
"""向 AutomationBench 5113 Golden 逐项注入本题专属错误，验证 Judge 拒绝坏答案。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
import zipfile
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path
from typing import Callable

from docx import Document
from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
XLSX = "I9_Compliance_Audit.xlsx"
DOCX = "I9_Compliance_Control_Brief.docx"
PPTX = "I9_Compliance_Review.pptx"
MBOX = "I9_Notification_Drafts.mbox"
UPDATES = "I9_Tracker_Update_Drafts.json"
SLACK = "I9_Slack_Control.json"
EXECUTION = "I9_Execution_Control.json"
FILES = (XLSX, DOCX, PPTX, MBOX, UPDATES, SLACK, EXECUTION)
Mutator = Callable[[Path], None]
Case = tuple[str, Mutator, str]


def load_judge():
    spec = importlib.util.spec_from_file_location("i9_5113_judge", ROOT / "tests" / "judge.py")
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入本题 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    module.render = lambda path: ({".xlsx": 9, ".docx": 3, ".pptx": 3}[path.suffix], "render-ok")
    return module


def rewrite_zip(path: Path, transform: Callable[[str, bytes], bytes | None]) -> None:
    temporary = path.with_suffix(path.suffix + ".negative")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = transform(info.filename, source.read(info.filename))
            if data is not None:
                target.writestr(info, data)
    temporary.replace(path)


def sheet_xml_path(archive: zipfile.ZipFile, sheet_name: str) -> str:
    workbook = etree.fromstring(archive.read("xl/workbook.xml"))
    relationships = etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    namespaces = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main", "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships", "pr": "http://schemas.openxmlformats.org/package/2006/relationships"}
    relationship_id = workbook.xpath(f'//m:sheet[@name="{sheet_name}"]/@r:id', namespaces=namespaces)[0]
    target = relationships.xpath(f'//pr:Relationship[@Id="{relationship_id}"]/@Target', namespaces=namespaces)[0]
    return target.lstrip("/") if target.startswith("/") else "xl/" + target.lstrip("/")


def edit_cell(path: Path, sheet_name: str, coordinate: str, *, value=None, remove_formula: bool = False) -> None:
    with zipfile.ZipFile(path, "r") as archive:
        target_sheet = sheet_xml_path(archive, sheet_name)
    changed = False
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

    def transform(name: str, data: bytes) -> bytes:
        nonlocal changed
        if name != target_sheet:
            return data
        root = etree.fromstring(data)
        cells = root.xpath(f'//m:c[@r="{coordinate}"]', namespaces={"m": namespace})
        if not cells:
            raise ValueError(f"{sheet_name}!{coordinate} 不存在")
        cell = cells[0]
        if remove_formula:
            formula = cell.find(f"{{{namespace}}}f")
            if formula is None:
                raise ValueError(f"{sheet_name}!{coordinate} 没有公式")
            cell.remove(formula)
        if value is not None:
            inline = cell.find(f"{{{namespace}}}is")
            if inline is not None:
                cell.remove(inline)
            node = cell.find(f"{{{namespace}}}v")
            if node is None:
                node = etree.SubElement(cell, f"{{{namespace}}}v")
            if isinstance(value, bool):
                node.text = "1" if value else "0"
                cell.set("t", "b")
            elif isinstance(value, int):
                node.text = str(value)
                cell.attrib.pop("t", None)
            else:
                node.text = str(value)
                cell.set("t", "str")
        changed = True
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    rewrite_zip(path, transform)
    if not changed:
        raise ValueError("目标单元格未修改")


def openpyxl_mutation(path: Path, mutator: Callable) -> None:
    workbook = load_workbook(path)
    mutator(workbook)
    workbook.save(path)


def mutate_json(path: Path, mutator: Callable[[dict], None]) -> None:
    payload = json.loads(path.read_text(encoding="utf-8"))
    mutator(payload)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def decode_message(message) -> str:
    if message.is_multipart():
        return "\n".join((part.get_payload(decode=True) or b"").decode(part.get_content_charset() or "utf-8", errors="replace") for part in message.walk() if part.get_content_type() == "text/plain").strip()
    return (message.get_payload(decode=True) or b"").decode(message.get_content_charset() or "utf-8", errors="replace").strip()


def mbox_states(path: Path) -> list[dict]:
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    return [{
        "from": str(message.get("From", "")), "to": str(message.get("To", "")), "cc": str(message.get("Cc", "")), "bcc": str(message.get("Bcc", "")),
        "subject": str(message.get("Subject", "")), "date": str(message.get("Date", "")), "message_id": str(message.get("Message-ID", "")),
        "status": str(message.get("X-EvoCowork-Status", "")), "source_row": str(message.get("X-EvoCowork-Source-Row", "")),
        "route": str(message.get("X-EvoCowork-Route", "")), "body": decode_message(message),
    } for message in messages]


def write_mbox_states(path: Path, states: list[dict]) -> None:
    chunks: list[bytes] = []
    for order, state in enumerate(states, start=1):
        message = EmailMessage(policy=SMTP)
        for key, header in (("from", "From"), ("to", "To"), ("cc", "Cc"), ("bcc", "Bcc"), ("subject", "Subject"), ("date", "Date"), ("message_id", "Message-ID"), ("status", "X-EvoCowork-Status"), ("source_row", "X-EvoCowork-Source-Row"), ("route", "X-EvoCowork-Route")):
            if state.get(key):
                message[header] = state[key]
        message.set_content(state["body"].rstrip() + "\n", charset="utf-8")
        buffer = BytesIO()
        BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
        chunks.extend([f"From MAILER-DAEMON Wed Apr 15 09:00:{order:02d} 2026\n".encode("ascii"), buffer.getvalue().replace(b"\r\n", b"\n"), b"\n"])
    path.write_bytes(b"".join(chunks))


def mutate_message(path: Path, index: int, mutator: Callable[[dict], None]) -> None:
    states = mbox_states(path)
    mutator(states[index])
    write_mbox_states(path, states)


def replace_docx_text(path: Path, old: str, new: str, *, row_token: str | None = None) -> None:
    document = Document(path)
    changed = False
    if row_token is None:
        for paragraph in document.paragraphs:
            if old in paragraph.text:
                paragraph.text = paragraph.text.replace(old, new, 1)
                changed = True
                break
    if not changed:
        for table in document.tables:
            for row in table.rows:
                if row_token and row_token not in " | ".join(cell.text for cell in row.cells):
                    continue
                for cell in row.cells:
                    if old in cell.text:
                        cell.text = cell.text.replace(old, new, 1)
                        changed = True
                        break
                if changed:
                    break
            if changed:
                break
    if not changed:
        raise ValueError(f"Word 找不到 {old!r}")
    document.save(path)


def replace_pptx_text(path: Path, old: str, new: str) -> None:
    presentation = Presentation(path)
    changed = False
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and old in shape.text:
                shape.text = shape.text.replace(old, new, 1)
                changed = True
                break
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        if old in cell.text:
                            cell.text = cell.text.replace(old, new, 1)
                            changed = True
                            break
                    if changed:
                        break
            if changed:
                break
        if changed:
            break
    if not changed:
        raise ValueError(f"PPT 找不到 {old!r}")
    presentation.save(path)


def replace_pptx_row_text(path: Path, row_token: str, old: str, new: str) -> None:
    presentation = Presentation(path)
    changed = False
    for slide in presentation.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_table", False):
                continue
            for row in shape.table.rows:
                if row_token not in " | ".join(cell.text for cell in row.cells):
                    continue
                for cell in row.cells:
                    if cell.text == old:
                        cell.text = new
                        changed = True
                        break
                if changed:
                    break
            if changed:
                break
        if changed:
            break
    if not changed:
        raise ValueError(f"PPT 在 {row_token!r} 行找不到精确单元格 {old!r}")
    presentation.save(path)


def drop_ppt_slide(path: Path) -> None:
    presentation = Presentation(path)
    presentation.slides._sldIdLst.remove(presentation.slides._sldIdLst[-1])
    presentation.save(path)


def remove_path(root: Path, name: str) -> None:
    (root / name).unlink()


def corrupt_path(root: Path, name: str) -> None:
    (root / name).write_bytes(b"not-a-valid-artifact")


def build_cases() -> list[Case]:
    cases: list[Case] = []
    for name in FILES:
        cases.append((f"missing-{name}", lambda root, n=name: remove_path(root, n), "R001"))
    cases.append(("extra-file", lambda root: (root / "unexpected.txt").write_text("unexpected", encoding="utf-8"), "R001"))
    for name in FILES:
        cases.append((f"corrupt-{name}", lambda root, n=name: corrupt_path(root, n), "R001" if name != MBOX else "R013"))

    source_edits = [
        ("amara-row", "A2", 12), ("amara-name", "B2", "Amara Wrong"), ("amara-email", "C2", "wrong@company.example.com"), ("amara-manager", "D2", "wrong-manager@company.example.com"),
        ("boris-start", "E3", "2026-04-08"), ("boris-status", "F3", "Complete"), ("cassie-name", "B4", "Cassie Wrong"), ("cassie-status", "F4", "Incomplete"),
        ("dmitri-email", "C5", "dmitri.wrong@company.example.com"), ("dmitri-manager", "D5", "mgr-eng@company.example.com"), ("esme-row", "A6", 16), ("esme-start", "E6", "2026-04-10"),
        ("frankie-name", "B7", "Frankie Wrong"), ("frankie-start", "E7", "2026-03-26"), ("gina-email", "C8", "gina.wrong@company.example.com"), ("gina-manager", "D8", "mgr-sales@company.example.com"),
        ("hank-row", "A9", 19), ("hank-start", "E9", "2026-04-09"), ("hank-status", "F9", "Complete"), ("hank-notes", "G9", "waived"),
    ]
    for name, cell, value in source_edits:
        cases.append((f"source-{name}", lambda root, c=cell, v=value: edit_cell(root / XLSX, "Source_Hires", c, value=v), "R002"))

    context_edits = [
        ("policy-message-id", "Policy_Source", "A2", "wrong-id"), ("policy-from", "Policy_Source", "B2", "hr@company.example.com"),
        ("policy-to", "Policy_Source", "C2", "all@company.example.com"), ("policy-subject", "Policy_Source", "D2", "Wrong policy"),
        ("policy-date", "Policy_Source", "E2", "Tue, 14 Apr 2026 09:00:00 +0000"), ("policy-body", "Policy_Source", "F2", "Incomplete policy"),
        ("rule-deadline", "Policy_Rules", "C4", "Complete in 5 business days"), ("rule-reminder", "Policy_Rules", "C5", "Send Slack reminder"),
        ("rule-legal", "Policy_Rules", "C6", "Email manager only"), ("rule-retention", "Policy_Rules", "C7", "Retain one month"),
        ("rule-audit", "Policy_Rules", "C8", "No audits"), ("rule-order", "Policy_Rules", "A5", 9),
        ("evidence-amara", "Evidence_Register", "D2", "Amara wrong evidence"), ("evidence-hank", "Evidence_Register", "D9", "Hank omitted"),
        ("evidence-policy", "Evidence_Register", "D10", "Wrong policy body"), ("evidence-time", "Evidence_Register", "D11", "2026-04-16T09:00:00Z"),
        ("evidence-channel", "Evidence_Register", "D12", '{"id":"C_HR_5113","name":"hr-general","is_private":true}'), ("evidence-trace", "Evidence_Register", "E12", "wrong-source"),
    ]
    for name, sheet, cell, value in context_edits:
        cases.append((name, lambda root, s=sheet, c=cell, v=value: edit_cell(root / XLSX, s, c, value=v), "R003"))

    formula_cells = [
        ("source-row", "Compliance_Decisions", "A2"), ("source-name", "Compliance_Decisions", "B3"), ("source-email", "Compliance_Decisions", "C4"),
        ("source-manager", "Compliance_Decisions", "D5"), ("start-date", "Compliance_Decisions", "E6"), ("source-status", "Compliance_Decisions", "F7"),
        ("source-notes", "Compliance_Decisions", "G9"), ("deadline", "Compliance_Decisions", "H2"), ("elapsed", "Compliance_Decisions", "I3"),
        ("complete", "Compliance_Decisions", "J4"), ("within", "Compliance_Decisions", "K6"), ("reminder", "Compliance_Decisions", "L8"),
        ("legal", "Compliance_Decisions", "M9"), ("no-action", "Compliance_Decisions", "N4"), ("proposed", "Compliance_Decisions", "O5"),
        ("email-draft", "Compliance_Decisions", "P7"), ("slack-route", "Compliance_Decisions", "Q2"), ("trace", "Compliance_Decisions", "R9"),
        ("reminder-queue", "Reminder_Queue", "C2"), ("reminder-cc", "Reminder_Queue", "E3"), ("legal-queue", "Legal_Violation_Queue", "C5"),
        ("no-action-queue", "No_Action_Queue", "E2"), ("summary-count", "Control_Summary", "B5"), ("summary-side-effect", "Control_Summary", "B14"),
    ]
    for name, sheet, cell in formula_cells:
        cases.append((f"formula-{name}", lambda root, s=sheet, c=cell: edit_cell(root / XLSX, s, c, remove_formula=True), "R004"))

    decision_edits = [
        ("amara-deadline", "H2", "2026-04-16", "R005"), ("amara-elapsed", "I2", 2, "R005"), ("amara-no-reminder", "L2", "NO", "R005"), ("amara-legal", "M2", "YES", "R005"),
        ("boris-deadline", "H3", "2026-04-11", "R008"), ("boris-elapsed", "I3", 5, "R008"), ("boris-reminder", "L3", "YES", "R008"), ("boris-no-legal", "M3", "NO", "R008"),
        ("cassie-complete", "J4", "NO", "R012"), ("cassie-reminder", "L4", "YES", "R012"), ("cassie-email", "P4", "YES", "R012"), ("cassie-proposed", "O4", "Reminder Sent", "R012"),
        ("dmitri-deadline", "H5", "2026-04-07", "R009"), ("dmitri-elapsed", "I5", 9, "R009"), ("dmitri-reminder", "L5", "YES", "R009"), ("dmitri-status", "O5", "Reminder Sent", "R009"),
        ("esme-elapsed", "I6", 4, "R006"), ("esme-within", "K6", "NO", "R006"), ("esme-no-reminder", "L6", "NO", "R006"), ("esme-legal", "M6", "YES", "R006"),
        ("frankie-deadline", "H7", "2026-03-31", "R010"), ("frankie-elapsed", "I7", 14, "R010"), ("frankie-reminder", "L7", "YES", "R010"), ("frankie-no-legal", "M7", "NO", "R010"),
        ("gina-elapsed", "I8", 1, "R007"), ("gina-within", "K8", "NO", "R007"), ("gina-no-reminder", "L8", "NO", "R007"), ("gina-legal", "M8", "YES", "R007"),
        ("hank-deadline", "H9", "2026-04-14", "R011"), ("hank-elapsed", "I9", 4, "R011"), ("hank-reminder", "L9", "YES", "R011"), ("hank-no-legal", "M9", "NO", "R011"),
    ]
    for name, cell, value, criterion in decision_edits:
        cases.append((name, lambda root, c=cell, v=value: edit_cell(root / XLSX, "Compliance_Decisions", c, value=v), criterion))

    queue_edits = [
        ("reminder-amara-name", "Reminder_Queue", "C2", "Amara Wrong", "R013"), ("reminder-amara-email", "Reminder_Queue", "D2", "wrong@company.example.com", "R013"),
        ("reminder-amara-cc", "Reminder_Queue", "E2", "wrong-manager@company.example.com", "R013"), ("reminder-esme-order", "Reminder_Queue", "A3", 3, "R013"),
        ("reminder-esme-elapsed", "Reminder_Queue", "H3", 4, "R013"), ("reminder-gina-subject", "Reminder_Queue", "I4", "Wrong subject", "R013"),
        ("reminder-sent", "Reminder_Queue", "K2", True, "R013"), ("legal-boris-name", "Legal_Violation_Queue", "C2", "Boris Wrong", "R013"),
        ("legal-dmitri-date", "Legal_Violation_Queue", "D3", "2026-04-02", "R013"), ("legal-frankie-recipient", "Legal_Violation_Queue", "G4", "manager@company.example.com", "R013"),
        ("legal-hank-name", "Legal_Violation_Queue", "C5", "Hank Missing", "R013"), ("legal-hank-sent", "Legal_Violation_Queue", "J5", True, "R013"),
        ("no-action-cassie", "No_Action_Queue", "C2", "Cassie Wrong", "R012"), ("no-action-email", "No_Action_Queue", "F2", True, "R012"),
        ("summary-reminders", "Control_Summary", "B5", 2, "R004"), ("summary-violations", "Control_Summary", "B6", 3, "R004"),
        ("summary-email-sent", "Control_Summary", "B9", 1, "R004"), ("summary-applied", "Control_Summary", "B11", 1, "R004"),
    ]
    for name, sheet, cell, value, criterion in queue_edits:
        cases.append((name, lambda root, s=sheet, c=cell, v=value: edit_cell(root / XLSX, s, c, value=v), criterion))
    cases.extend([
        ("xlsx-native-table-removed", lambda root: openpyxl_mutation(root / XLSX, lambda wb: wb["Source_Hires"].tables.clear()), "R018"),
        ("xlsx-native-chart-removed", lambda root: openpyxl_mutation(root / XLSX, lambda wb: setattr(wb["Control_Summary"], "_charts", [])), "R018"),
        ("xlsx-conditional-formatting-removed", lambda root: openpyxl_mutation(root / XLSX, lambda wb: setattr(wb["Compliance_Decisions"], "conditional_formatting", type(wb["Compliance_Decisions"].conditional_formatting)())), "R018"),
    ])

    cases.extend([
        ("mbox-empty", lambda root: write_mbox_states(root / MBOX, []), "R013"),
        ("mbox-drop-amara", lambda root: write_mbox_states(root / MBOX, mbox_states(root / MBOX)[1:]), "R013"),
        ("mbox-drop-hank", lambda root: write_mbox_states(root / MBOX, mbox_states(root / MBOX)[:-1]), "R013"),
        ("mbox-duplicate", lambda root: write_mbox_states(root / MBOX, mbox_states(root / MBOX) + [mbox_states(root / MBOX)[-1]]), "R013"),
        ("mbox-reverse", lambda root: write_mbox_states(root / MBOX, list(reversed(mbox_states(root / MBOX)))), "R013"),
        ("mbox-swap", lambda root: write_mbox_states(root / MBOX, [mbox_states(root / MBOX)[1], mbox_states(root / MBOX)[0], *mbox_states(root / MBOX)[2:]]), "R005"),
    ])
    message_cases = [
        ("amara-to", 0, "to", "legal@company.example.com", "R005"), ("amara-cc", 0, "cc", "mgr-sales@company.example.com", "R005"), ("amara-bcc", 0, "bcc", "audit@company.example.com", "R005"),
        ("amara-subject", 0, "subject", "I-9 Notice", "R005"), ("amara-date", 0, "date", "Thu, 16 Apr 2026 09:00:00 +0000", "R005"), ("amara-status", 0, "status", "SENT", "R005"),
        ("boris-to", 1, "to", "boris.petrov@company.example.com", "R008"), ("boris-cc", 1, "cc", "mgr-sales@company.example.com", "R008"), ("boris-subject", 1, "subject", "I-9 Reminder", "R008"),
        ("boris-route", 1, "route", "REMINDER_WITH_MANAGER_CC", "R008"), ("dmitri-source", 2, "source_row", "6", "R009"), ("dmitri-message-id", 2, "message_id", "<wrong@evocowork.local>", "R009"),
        ("esme-to", 3, "to", "legal@company.example.com", "R006"), ("esme-cc", 3, "cc", "mgr-eng@company.example.com", "R006"), ("esme-status", 3, "status", "SENT", "R006"),
        ("frankie-subject", 4, "subject", "I-9 Compliance", "R010"), ("frankie-bcc", 4, "bcc", "audit@company.example.com", "R010"), ("gina-cc", 5, "cc", "mgr-sales@company.example.com", "R007"),
        ("gina-route", 5, "route", "VIOLATION_LEGAL_NOTIFICATION", "R007"), ("hank-to", 6, "to", "hank.morrison@company.example.com", "R011"), ("hank-subject", 6, "subject", "I-9 Reminder", "R011"),
    ]
    for name, index, field, value, criterion in message_cases:
        cases.append((f"mbox-{name}", lambda root, i=index, f=field, v=value: mutate_message(root / MBOX, i, lambda state: state.update({f: v})), criterion))
    body_cases = [
        ("amara-start", 0, "Start Date=2026-04-14", "Start Date=2026-04-13", "R005"), ("amara-elapsed", 0, "elapsed business days=1", "elapsed business days=2", "R005"),
        ("boris-name", 1, "Employee=Boris Petrov", "Employee=Boris Wrong", "R008"), ("boris-deadline", 1, "deadline=2026-04-10", "deadline=2026-04-11", "R008"),
        ("dmitri-start", 2, "Start Date=2026-04-01", "Start Date=2026-04-02", "R009"), ("esme-elapsed", 3, "elapsed business days=3", "elapsed business days=4", "R006"),
        ("frankie-sent", 4, "has not been sent", "has been sent", "R010"), ("gina-deadline", 5, "deadline=2026-04-20", "deadline=2026-04-19", "R007"),
        ("hank-name", 6, "Employee=Hank Morrison", "Employee=Hank Missing", "R011"), ("hank-elapsed", 6, "elapsed business days=5", "elapsed business days=4", "R011"),
    ]
    for name, index, old, new, criterion in body_cases:
        def mutate_body(root, i=index, o=old, n=new):
            def change(state):
                if o not in state["body"]:
                    raise ValueError(f"邮件正文找不到 {o!r}")
                state["body"] = state["body"].replace(o, n, 1)
            mutate_message(root / MBOX, i, change)
        cases.append((f"mbox-body-{name}", mutate_body, criterion))

    update_cases = [
        ("task-id", lambda p: p.update({"task_id": "wrong"})), ("status", lambda p: p.update({"status": "APPLIED"})), ("frozen-time", lambda p: p.update({"frozen_time": "2026-04-16T09:00:00Z"})),
        ("method", lambda p: p.update({"business_day_method": "calendar days"})), ("draft-count", lambda p: p.update({"draft_count": 6})), ("applied-count", lambda p: p.update({"applied_count": 1})),
        ("drop-amara", lambda p: p["drafts"].pop(0)), ("drop-hank", lambda p: p["drafts"].pop()), ("reverse", lambda p: p["drafts"].reverse()),
        ("amara-email-route", lambda p: p["drafts"][0].update({"route": "VIOLATION_LEGAL_NOTIFICATION"})), ("amara-status", lambda p: p["drafts"][0].update({"proposed_status": "Complete"})),
        ("boris-date", lambda p: p["drafts"][1].update({"start_date": "2026-04-08"})), ("boris-elapsed", lambda p: p["drafts"][1].update({"elapsed_business_days": 5})),
        ("dmitri-deadline", lambda p: p["drafts"][2].update({"deadline": "2026-04-07"})), ("esme-route", lambda p: p["drafts"][3].update({"route": "VIOLATION_LEGAL_NOTIFICATION"})),
        ("frankie-applied", lambda p: p["drafts"][4].update({"applied": True, "state": "APPLIED"})), ("gina-row", lambda p: p["drafts"][5].update({"source_row": 9})),
        ("hank-name", lambda p: p["drafts"][6].update({"employee": "Hank Missing"})), ("hank-trace", lambda p: p["drafts"][6].update({"source_trace": []})),
        ("add-cassie", lambda p: p["drafts"].append({"employee": "Cassie Lin"})),
    ]
    for name, mutator in update_cases:
        cases.append((f"updates-{name}", lambda root, m=mutator: mutate_json(root / UPDATES, m), "R013" if name != "add-cassie" else "R012"))

    slack_cases = [
        ("task-id", lambda p: p.update({"task_id": "wrong"})), ("channel-id", lambda p: p["source_channel"].update({"id": "C_OTHER"})), ("channel-private", lambda p: p["source_channel"].update({"is_private": True})),
        ("decision", lambda p: p.update({"decision": "POST"})), ("reason", lambda p: p.update({"reason": "Slack authorized"})), ("draft", lambda p: p.update({"draft_count": 1})),
        ("posted-count", lambda p: p.update({"posted_count": 1})), ("posted", lambda p: p.update({"posted": True})), ("employee-leak", lambda p: p["employee_names_in_slack"].append("Hank Morrison")),
        ("status-leak", lambda p: p["i9_statuses_in_slack"].append("Incomplete")), ("message", lambda p: p["messages_created"].append({"text": "I-9"})), ("trace", lambda p: p.update({"source_trace": []})),
    ]
    for name, mutator in slack_cases:
        cases.append((f"slack-{name}", lambda root, m=mutator: mutate_json(root / SLACK, m), "R014"))

    execution_cases = [
        ("task-id", lambda p: p.update({"task_id": "wrong"})), ("status", lambda p: p.update({"status": "EXECUTED"})), ("side-effect", lambda p: p.update({"real_side_effect_count": 1})),
        ("email-reminders", lambda p: p["email"].update({"reminder_drafts": 2})), ("email-legal", lambda p: p["email"].update({"legal_violation_drafts": 3})),
        ("email-total", lambda p: p["email"].update({"total_drafts": 6})), ("email-sent", lambda p: p["email"].update({"sent_count": 1, "state": "SENT"})),
        ("tracker-reminders", lambda p: p["tracker_updates"].update({"reminder_status_drafts": 2})), ("tracker-legal", lambda p: p["tracker_updates"].update({"violation_status_drafts": 3})),
        ("tracker-total", lambda p: p["tracker_updates"].update({"total_drafts": 6})), ("tracker-applied", lambda p: p["tracker_updates"].update({"applied_count": 1, "state": "APPLIED"})),
        ("cassie-name", lambda p: p["no_action"].update({"employee": "Hank Morrison"})), ("cassie-email", lambda p: p["no_action"].update({"email_draft_count": 1})),
        ("cassie-update", lambda p: p["no_action"].update({"tracker_update_draft_count": 1})), ("slack-draft", lambda p: p["slack"].update({"draft_count": 1})),
        ("slack-post", lambda p: p["slack"].update({"posted_count": 1, "state": "POSTED"})), ("channel", lambda p: p["slack"].update({"channel": "C_OTHER"})),
        ("trace", lambda p: p.update({"source_trace": []})),
    ]
    for name, mutator in execution_cases:
        cases.append((f"execution-{name}", lambda root, m=mutator: mutate_json(root / EXECUTION, m), "R014"))

    word_cases = [
        ("summary-reminder", "Reminder", "Wrong reminder", "Reminder", "R016"), ("summary-violation-count", "4", "3", "Violation", "R016"),
        ("amara-elapsed", "1", "2", "Amara Obi", "R005"), ("boris-deadline", "2026-04-10", "2026-04-11", "Boris Petrov", "R008"),
        ("cassie-source", "Complete", "Incomplete", "Cassie Lin", "R012"), ("dmitri-status", "Violation - Legal Notified", "Reminder Sent", "Dmitri Volkov", "R009"),
        ("esme-elapsed", "3", "4", "Esme Nakamura", "R006"), ("frankie-name", "Frankie Reeves", "Frankie Missing", "Frankie Reeves", "R010"),
        ("gina-manager", "mgr-mkt@company.example.com", "mgr-sales@company.example.com", "Gina Torres", "R007"), ("hank-name", "Hank Morrison", "Hank Missing", "Hank Morrison", "R011"),
        ("policy-days", "3 business days", "5 business days", "Employer deadline", "R016"), ("email-sent", "DRAFT_NOT_SENT", "SENT", "Emails", "R016"),
        ("tracker-applied", "DRAFT_NOT_APPLIED", "APPLIED", "Tracker updates", "R016"), ("slack-posted", "NOT_AUTHORIZED_CHANNEL", "POSTED", "Slack", "R016"),
        ("side-effect", "OFFLINE_ONLY", "EXECUTED", "Real side effects", "R016"),
    ]
    for name, old, new, row_token, criterion in word_cases:
        cases.append((f"word-{name}", lambda root, o=old, n=new, t=row_token: replace_docx_text(root / DOCX, o, n, row_token=t), criterion))
    cases.extend([
        ("word-method", lambda root: replace_docx_text(root / DOCX, "Monday through Friday", "Every calendar day"), "R016"),
        ("word-hank-assertion", lambda root: replace_docx_text(root / DOCX, "upstream assertions do not name him", "upstream assertions name him"), "R016"),
        ("word-real-side-effects", lambda root: replace_docx_text(root / DOCX, "real side effects=0", "real side effects=1"), "R016"),
    ])

    ppt_cases = [
        ("summary-hires", "Eight reviewed", "Seven reviewed", "R015"), ("summary-reminders", "three reminders", "two reminders", "R015"),
        ("summary-violations", "four violations", "three violations", "R015"), ("summary-side-effects", "zero side effects", "one side effect", "R015"),
        ("amara-name", "Amara Obi", "Amara Wrong", "R005"),
        ("dmitri-name", "Dmitri Volkov", "Dmitri Wrong", "R009"),
        ("frankie-name", "Frankie Reeves", "Frankie Missing", "R010"),
        ("hank-name", "Hank Morrison", "Hank Missing", "R011"),
        ("sent", "no email sent", "email sent", "R015"), ("applied", "no tracker update applied", "tracker update applied", "R015"),
        ("slack", "no Slack message", "Slack message posted", "R015"), ("side-effect", "Real side effects=0", "Real side effects=1", "R015"),
    ]
    for name, old, new, criterion in ppt_cases:
        cases.append((f"ppt-{name}", lambda root, o=old, n=new: replace_pptx_text(root / PPTX, o, n), criterion))
    cases.extend([
        ("ppt-boris-elapsed", lambda root: replace_pptx_row_text(root / PPTX, "Boris Petrov", "6", "5"), "R008"),
        ("ppt-cassie-status", lambda root: replace_pptx_row_text(root / PPTX, "Cassie Lin", "Complete", "Incomplete"), "R012"),
        ("ppt-esme-deadline", lambda root: replace_pptx_row_text(root / PPTX, "Esme Nakamura", "2026-04-15", "2026-04-16"), "R006"),
        ("ppt-gina-route", lambda root: replace_pptx_row_text(root / PPTX, "Gina Torres", "REMINDER", "LEGAL"), "R007"),
        ("ppt-drop-slide", lambda root: drop_ppt_slide(root / PPTX), "R015"),
        ("ppt-hank-policy", lambda root: replace_pptx_text(root / PPTX, "Hank is included by full-roster policy evaluation", "Hank omitted"), "R015"),
        ("ppt-cassie-protection", lambda root: replace_pptx_text(root / PPTX, "Cassie is protected by Complete source status", "Cassie receives reminder"), "R015"),
    ])
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--json-out", type=Path)
    args = parser.parse_args()
    judge = load_judge()
    cases = build_cases()
    results: list[dict] = []
    failures: list[str] = []
    for name, mutator, expected_failure in cases:
        with tempfile.TemporaryDirectory(prefix=f"negative-5113-{name}-") as temporary:
            artifact_root = Path(temporary) / "artifacts"
            shutil.copytree(GOLDEN, artifact_root)
            try:
                mutator(artifact_root)
                result = judge.evaluate(artifact_root)
                detected = not result.get("criteria", {}).get(expected_failure, {}).get("passed", True)
                rejected = not result.get("passed", False) and float(result.get("reward", 1.0)) < 0.8
                results.append({"case": name, "expected_failure": expected_failure, "detected": detected, "rejected": rejected, "reward": result.get("reward"), "hurdle_failures": result.get("hurdle_failures", [])})
                if not detected or not rejected:
                    failures.append(f"{name}: expected={expected_failure}, detected={detected}, reward={result.get('reward')}")
            except Exception as exc:
                failures.append(f"{name}: mutation/evaluation error: {type(exc).__name__}: {exc}")
    payload = {"task_id": "cowork-evo/automationbench-5113-i9-compliance-control-package", "case_count": len(results), "declared_case_count": len(cases), "all_rejected": not failures and len(results) == len(cases), "max_reward": max((float(item["reward"]) for item in results), default=None), "failures": failures, "results": results}
    if args.json_out:
        args.json_out.parent.mkdir(parents=True, exist_ok=True)
        args.json_out.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    if failures or len(results) != len(cases):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
