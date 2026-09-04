#!/usr/bin/env python3
"""AutomationBench 5114 的 v2 连续评分入口；原严格口径保存在 judge_v1.py。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import mailbox
import sys
import tempfile
from email.utils import getaddresses
from pathlib import Path
from typing import Any, Callable

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


TASK_ID = "cowork-evo/automationbench-5114-termination-checklist-governance-package"
RUBRIC_VERSION = "automationbench-5114-termination-v0.2"
JUDGE_VERSION = "automationbench-5114-criterion-partial-v2.0"
XLSX = "Termination_Checklist_Audit.xlsx"
DOCX = "Termination_Checklist_Brief.docx"
PPTX = "Termination_Checklist_Review.pptx"
MBOX = "Termination_Communication_Drafts.mbox"
SLACK = "HR_Ops_Termination_Slack_Drafts.json"
UPDATES = "Termination_Tracker_Update_Drafts.json"
EXECUTION = "Termination_Execution_Control.json"
REQUIRED = (XLSX, DOCX, PPTX, MBOX, SLACK, UPDATES, EXECUTION)
POSITIVE_WEIGHTS = {
    "R001": 6, "R002": 8, "R003": 10, "R004": 10,
    "R005": 10, "R006": 8, "R007": 9, "R008": 8,
    "R009": 8, "R010": 8, "R011": 7, "R012": 8,
}
PENALTY_WEIGHTS = {"P001": -30, "P002": -30, "P003": -25}
PROCESS = [
    (2, "Aaron Wells", "aaron.wells@company.example.com", "2026-04-18", "5 days"),
    (3, "Bianca Morales", "bianca.morales@company.example.com", "2026-04-16", "PTO balance unknown"),
    (4, "Chad Eriksson", "chad.eriksson@company.example.com", "2026-04-20", "12 days"),
    (6, "Ethan Park", "ethan.park@company.example.com", "2026-04-22", "PTO balance unknown"),
    (7, "Fiona Strand", "fiona.strand@company.example.com", "2026-04-25", "8 days"),
    (8, "George Tan", "george.tan@company.example.com", "2026-04-17", "0 days"),
    (9, "Holly Ng", "holly.ng@company.example.com", "2026-04-30", "PTO balance unknown"),
]
SKIPS = [
    (5, "Dina Petrova", "Processed", "Already completed", "SKIP_ALREADY_PROCESSED", "Source status is Processed"),
    (10, "Isaac Reeves", "Pending", "Rescinded - employee accepted counteroffer, staying", "SKIP_RESCINDED", "Notes say employee accepted counteroffer and is staying"),
]


def _load_v1():
    path = Path(__file__).with_name("judge_v1.py")
    spec = importlib.util.spec_from_file_location("automationbench_5114_judge_v1", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v1 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


V1 = _load_v1()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _norm(value: Any) -> Any:
    if value is None:
        return ""
    if hasattr(value, "date") and not isinstance(value, str):
        return value.date().isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _rows(sheet, start: int = 2, end: int | None = None, columns: int | None = None) -> list[tuple[Any, ...]]:
    end = end or sheet.max_row
    columns = columns or sheet.max_column
    return [tuple(_norm(sheet.cell(row, column).value) for column in range(1, columns + 1)) for row in range(start, end + 1)]


def _messages(path: Path) -> list:
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    return messages


def _decode(message) -> str:
    if message.is_multipart():
        return "\n".join(
            (part.get_payload(decode=True) or b"").decode(part.get_content_charset() or "utf-8", errors="replace")
            for part in message.walk() if part.get_content_type() == "text/plain"
        ).strip()
    return (message.get_payload(decode=True) or b"").decode(message.get_content_charset() or "utf-8", errors="replace").strip()


def _addresses(message, header: str) -> set[str]:
    return {address.casefold() for _, address in getaddresses(message.get_all(header, [])) if address}


def _core(artifact: Path) -> tuple[bool, list[str]]:
    errors: list[str] = []
    for name in REQUIRED:
        path = artifact / name
        if not path.is_file() or path.stat().st_size == 0:
            errors.append(f"核心交付缺失或为空：{name}")
            continue
        try:
            if name == XLSX:
                load_workbook(path, read_only=True, data_only=False).close()
            elif name == DOCX:
                Document(path)
            elif name == PPTX:
                if len(Presentation(path).slides) == 0:
                    raise ValueError("PPT 无页面")
            elif name == MBOX:
                messages = _messages(path)
                if not messages or not all(message.get("To") and message.get("Subject") for message in messages):
                    raise ValueError("MBOX 无有效终止沟通草稿")
            else:
                payload = json.loads(path.read_text(encoding="utf-8"))
                if not isinstance(payload, dict):
                    raise ValueError("JSON 顶层不是对象")
        except Exception as exc:
            errors.append(f"核心交付无法解析：{name}: {exc}")
    return not errors, errors


def _criterion(key: str, score: float, evidence: str, earned: float = 0, total: float = 1) -> dict[str, Any]:
    return {
        "criterion_id": key,
        "score": round(max(0.0, min(1.0, score)), 6),
        "evidence": evidence,
        "method": "deterministic",
        "raw": {"earned_units": round(earned, 6), "total_units": round(total, 6)},
    }


def _fatal(errors: list[str], hashes: list[str]) -> dict[str, Any]:
    return {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes or ["0" * 64],
        "criterion_results": [_criterion(key, 0, "；".join(errors), 0, 0) for key in (*POSITIVE_WEIGHTS, *PENALTY_WEIGHTS)],
        "reward": 0.0, "passed": False, "hurdle_failures": ["R001"], "penalties": [],
    }


def _ratio(findings: list[tuple[bool, str]]) -> tuple[float, list[str], int, int]:
    earned = sum(bool(ok) for ok, _ in findings)
    return earned / len(findings), [label for ok, label in findings if not ok], earned, len(findings)


def _legacy_exact(artifact: Path) -> dict[str, Any]:
    with tempfile.TemporaryDirectory(prefix="termination-5114-v1-") as temporary:
        exact = Path(temporary)
        for name in REQUIRED:
            (exact / name).symlink_to((artifact / name).resolve())
        return V1.evaluate(exact)


def _queue_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    book = load_workbook(artifact / XLSX, data_only=True)
    process = _rows(book["Process_Queue"], 2, 8, 5)
    it_rows = _rows(book["IT_Revocation_Drafts"], 2, 8, 6)
    payroll_rows = _rows(book["Payroll_Request_Drafts"], 2, 8, 6)
    tracker_rows = _rows(book["Tracker_Update_Drafts"], 2, 8, 5)
    book.close()
    findings: list[tuple[bool, str]] = []
    for observed, (source_row, employee, _email, last_day, pto) in zip(process, PROCESS):
        findings.append((observed == (source_row, employee, last_day, pto, "PROCESS_TERMINATION_CHECKLIST"), f"{employee} 处理队列事实或 PTO 错误"))
    for observed, (_source_row, employee, _email, last_day, pto) in zip(it_rows, PROCESS):
        findings.append((observed == (employee, "it-ops@company.example.com", last_day, pto, "DRAFT_NOT_SENT", "Access revocation request"), f"{employee} IT 请求队列错误"))
    for observed, (_source_row, employee, _email, last_day, pto) in zip(payroll_rows, PROCESS):
        findings.append((observed == (employee, "payroll@company.example.com", last_day, pto, "DRAFT_NOT_SENT", "Final check calculation request"), f"{employee} Payroll 请求队列错误"))
    for observed, (source_row, employee, _email, _last_day, _pto) in zip(tracker_rows, PROCESS):
        findings.append((observed == (source_row, employee, "Pending", "Processed", 0), f"{employee} Tracker 更新队列错误"))
    findings.extend([
        (len(process) == 7, "处理队列不是恰好七人"),
        (len(it_rows) == len(payroll_rows) == len(tracker_rows) == 7, "IT、Payroll 或 Tracker 队列不是各七条"),
    ])
    return _ratio(findings)


def _skip_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    book = load_workbook(artifact / XLSX, data_only=True)
    observed = _rows(book["Skip_Register"], 2, 3, 6)
    book.close()
    updates = json.loads((artifact / UPDATES).read_text(encoding="utf-8"))
    skipped = updates.get("skipped") if isinstance(updates.get("skipped"), list) else []
    findings = [
        (observed[0] == SKIPS[0] if len(observed) > 0 else False, "Dina 已处理跳过记录错误"),
        (observed[1] == SKIPS[1] if len(observed) > 1 else False, "Isaac 撤回跳过记录错误"),
        (skipped == [{"source_row": 5, "employee": "Dina Petrova", "decision": "SKIP_ALREADY_PROCESSED"}, {"source_row": 10, "employee": "Isaac Reeves", "decision": "SKIP_RESCINDED"}], "Update JSON 两条跳过记录错误"),
    ]
    return _ratio(findings)


def _email_body(kind: str, employee: str, last_day: str, pto: str) -> str:
    if kind == "it":
        return f"Please schedule access revocation for {employee} after the last day {last_day}. This is a request only; no access has been revoked. DRAFT_NOT_SENT."
    if kind == "payroll":
        return f"Please calculate the final paycheck for {employee}. Last day: {last_day}. PTO: {pto}. HR has not issued any payment. DRAFT_NOT_SENT."
    return f"Hello {employee},\n\nYour recorded last day is {last_day}. Please follow the standard offboarding instructions from HR. DRAFT_NOT_SENT."


def _communication_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    messages = _messages(artifact / MBOX)
    findings: list[tuple[bool, str]] = [(len(messages) == 21, "MBOX 不是恰好 21 封草稿")]
    kinds = (
        ("it", "IT_ACCESS_REVOCATION_REQUEST", "Access revocation request - ", "it-ops@company.example.com"),
        ("payroll", "PAYROLL_FINAL_CHECK_REQUEST", "Final check calculation request - ", "payroll@company.example.com"),
        ("employee", "EMPLOYEE_FAREWELL", "Offboarding instructions and farewell", None),
    )
    for employee_index, (_source_row, employee, email, last_day, pto) in enumerate(PROCESS):
        for kind_index, (kind, message_type, subject, recipient) in enumerate(kinds):
            position = employee_index * 3 + kind_index
            message = messages[position] if position < len(messages) else None
            if message is None:
                findings.extend([(False, f"缺少 {employee} {kind} 草稿") for _ in range(6)])
                continue
            wanted_subject = subject + employee if kind != "employee" else subject
            findings.extend([
                (_addresses(message, "From") == {"hr-ops@company.example.com"}, f"{employee} {kind} 发件人错误"),
                (_addresses(message, "To") == {recipient or email} and not _addresses(message, "Cc") and not _addresses(message, "Bcc"), f"{employee} {kind} 收件人或 Cc/Bcc 错误"),
                (str(message.get("Subject", "")) == wanted_subject, f"{employee} {kind} 主题错误"),
                (str(message.get("Date", "")) == "Wed, 15 Apr 2026 09:00:00 +0000" and str(message.get("Message-ID", "")) == f"<termination-{kind}-5114-{position + 1:02d}@draft.invalid>", f"{employee} {kind} 时间或 Message-ID 错误"),
                (str(message.get("X-EvoCowork-State", "")) == "DRAFT_NOT_SENT" and str(message.get("X-EvoCowork-Type", "")) == message_type, f"{employee} {kind} 状态或类型错误"),
                (_decode(message) == _email_body(kind, employee, last_day, pto), f"{employee} {kind} 姓名、日期、PTO 或未执行正文错误"),
            ])
    all_text = "\n".join(_decode(message) for message in messages)
    findings.append(("Dina Petrova" not in all_text and "Isaac Reeves" not in all_text, "跳过人员出现在 MBOX"))
    return _ratio(findings)


def _workflow_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    slack = json.loads((artifact / SLACK).read_text(encoding="utf-8"))
    updates = json.loads((artifact / UPDATES).read_text(encoding="utf-8"))
    slack_rows = slack.get("drafts") if isinstance(slack.get("drafts"), list) else []
    update_rows = updates.get("rows") if isinstance(updates.get("rows"), list) else []
    findings: list[tuple[bool, str]] = [
        (slack.get("task_id") == "automationbench-5114" and slack.get("channel") == "hr-ops", "Slack 任务身份或频道错误"),
        (slack.get("draft_count") == 7 and slack.get("posted_count") == 0 and slack.get("state") == "DRAFT_NOT_POSTED" and len(slack_rows) == 7, "Slack 草稿/发布计数或状态错误"),
        (updates.get("task_id") == "automationbench-5114" and updates.get("state") == "DRAFT_NOT_APPLIED", "Tracker Update 身份或状态错误"),
        (updates.get("draft_count") == 7 and updates.get("applied_count") == 0 and len(update_rows) == 7, "Tracker 草稿/应用计数错误"),
    ]
    for position, (source_row, employee, _email, last_day, _pto) in enumerate(PROCESS):
        slack_item = slack_rows[position] if position < len(slack_rows) and isinstance(slack_rows[position], dict) else {}
        update_item = update_rows[position] if position < len(update_rows) and isinstance(update_rows[position], dict) else {}
        findings.extend([
            (slack_item == {"employee": employee, "last_day": last_day, "text": f"Termination checklist draft prepared for {employee}; last day {last_day}. No action executed."}, f"{employee} Slack 草稿错误"),
            (update_item == {"source_row": source_row, "employee": employee, "from": "Pending", "to": "Processed"}, f"{employee} Tracker 更新草稿错误"),
        ])
    return _ratio(findings)


def _execution_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    payload = json.loads((artifact / EXECUTION).read_text(encoding="utf-8"))
    expected = {
        "task_id": "automationbench-5114", "reviewed": 9, "process_checklist": 7,
        "already_processed_skipped": 1, "rescinded_skipped": 1,
        "it_drafts": 7, "payroll_drafts": 7, "employee_drafts": 7,
        "email_drafts": 21, "emails_sent": 0, "slack_drafts": 7, "slack_posts": 0,
        "tracker_update_drafts": 7, "tracker_updates_applied": 0,
        "access_revocations_executed": 0, "final_paychecks_issued": 0, "real_side_effects": 0,
    }
    return _ratio([(payload.get(key) == value, f"Execution {key} 错误") for key, value in expected.items()])


def _word_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    doc = Document(artifact / DOCX)
    rows = [tuple(cell.text for cell in row.cells) for table in doc.tables for row in table.rows]
    text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(" | ".join(row) for row in rows)
    findings = [
        (len(doc.sections) == 3 and len(doc.tables) == 3, "Word 不是三页三表"),
        (("PROCESS_TERMINATION_CHECKLIST", "7", "Aaron; Bianca; Chad; Ethan; Fiona; George; Holly") in rows, "Word 七人处理汇总错误"),
        (("SKIP_ALREADY_PROCESSED", "1", "Dina") in rows, "Word Dina 跳过汇总错误"),
        (("SKIP_RESCINDED", "1", "Isaac") in rows, "Word Isaac 跳过汇总错误"),
    ]
    for _source_row, employee, _email, last_day, pto in PROCESS:
        findings.append(((employee, last_day, pto) in rows, f"Word {employee} Last day 或 PTO 行错误"))
    for wanted in [
        ("IT revocation request", "7", "0", "Not sent"),
        ("Payroll calculation request", "7", "0", "Not sent"),
        ("Employee farewell", "7", "0", "Not sent"),
        ("#hr-ops notice", "7", "0", "Not posted"),
        ("Tracker update", "7", "0", "Not applied"),
        ("Real side effects", "0", "0", "None"),
    ]:
        findings.append((wanted in rows, f"Word {wanted[0]} 控制行错误"))
    for token in ("Nine rows were reviewed", "Seven pending terminations", "Dina Petrova", "Isaac Reeves", "PTO balance unknown", "does not issue final pay", "No access was revoked", "no paycheck was issued", "no email was sent", "no source row was changed"):
        findings.append((token.casefold() in text.casefold(), f"Word 缺少 {token}"))
    return _ratio(findings)


def _ppt_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    deck = Presentation(artifact / PPTX)
    slide_texts: list[str] = []
    table_rows: list[list[tuple[str, ...]]] = []
    pictures = 0
    table_counts: list[int] = []
    for slide in deck.slides:
        parts: list[str] = []
        rows: list[tuple[str, ...]] = []
        count = 0
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                parts.append(shape.text.strip())
            if getattr(shape, "has_table", False):
                count += 1
                rows.extend(tuple(cell.text for cell in row.cells) for row in shape.table.rows)
        slide_texts.append("\n".join(parts))
        table_rows.append(rows)
        table_counts.append(count)
    findings: list[tuple[bool, str]] = [
        (len(deck.slides) == 3 and table_counts == [1, 1, 1] and pictures == 0, "PPT 不是三页、每页一张原生表或含图片替代"),
        (len(slide_texts) >= 1 and all(token in slide_texts[0] for token in ("Termination checklist scope", "9 reviewed", "7 process", "1 already processed", "1 rescinded")), "PPT 第1页摘要错误"),
        (len(table_rows) >= 1 and ("PROCESS", "7", "Aaron; Bianca; Chad; Ethan; Fiona; George; Holly") in table_rows[0], "PPT 第1页七人处理行错误"),
        (len(table_rows) >= 1 and ("SKIP PROCESSED", "1", "Dina") in table_rows[0], "PPT 第1页 Dina 跳过行错误"),
        (len(table_rows) >= 1 and ("SKIP RESCINDED", "1", "Isaac") in table_rows[0], "PPT 第1页 Isaac 跳过行错误"),
        (len(slide_texts) >= 2 and "HR requests calculation; unknown balances are stated explicitly" in slide_texts[1], "PPT 第2页 Payroll 边界说明错误"),
    ]
    if len(table_rows) >= 2:
        for _source_row, employee, _email, last_day, pto in PROCESS:
            findings.append(((employee, last_day, pto) in table_rows[1], f"PPT 第2页 {employee} 日期或 PTO 行错误"))
    if len(table_rows) >= 3:
        for wanted in [
            ("Emails", "21", "0", "DRAFT_NOT_SENT"),
            ("Slack notices", "7", "0", "DRAFT_NOT_POSTED"),
            ("Tracker updates", "7", "0", "NOT_APPLIED"),
            ("Revocations / paychecks", "0", "0", "NONE"),
        ]:
            findings.append((wanted in table_rows[2], f"PPT 第3页 {wanted[0]} 控制行错误"))
    findings.append((len(slide_texts) >= 3 and "Nothing is executed" in slide_texts[2] and "remain offline drafts" in slide_texts[2], "PPT 第3页零执行标题错误"))
    return _ratio(findings)


def _accuracy_severity(artifact: Path) -> tuple[float, list[str], float]:
    source_book = load_workbook(V1.inp() / "Termination_Tracker.xlsx", data_only=True)
    expected_source = _rows(source_book["Pending Terminations"], 2, 10, 8)
    source_book.close()
    book = load_workbook(artifact / XLSX, data_only=True)
    observed_source = _rows(book["Source_Tracker"], 2, 10, 8)
    process = _rows(book["Process_Queue"], 2, 8, 5)
    skips = _rows(book["Skip_Register"], 2, 3, 6)
    book.close()
    expected_process = [(row, name, day, pto, "PROCESS_TERMINATION_CHECKLIST") for row, name, _email, day, pto in PROCESS]
    failures: list[str] = []
    wrong = 0
    for index, (observed, expected) in enumerate(zip(observed_source, expected_source), start=2):
        if observed != expected:
            wrong += 1
            failures.append(f"Source_Tracker 第 {index} 行错误")
    for observed, expected in zip(process, expected_process):
        if observed != expected:
            wrong += 1
            failures.append(f"Process_Queue {expected[1]} 行错误")
    for observed, expected in zip(skips, SKIPS):
        if observed != expected:
            wrong += 1
            failures.append(f"Skip_Register {expected[1]} 行错误")
    total = len(expected_source) + len(expected_process) + len(SKIPS)
    return wrong / total, failures, total


def _boundary_severity(artifact: Path) -> tuple[float, list[str], float]:
    messages = _messages(artifact / MBOX)
    mail_text = "\n".join(_decode(message) for message in messages)
    slack = json.loads((artifact / SLACK).read_text(encoding="utf-8"))
    updates = json.loads((artifact / UPDATES).read_text(encoding="utf-8"))
    execution = json.loads((artifact / EXECUTION).read_text(encoding="utf-8"))
    action_names = mail_text + "\n" + "\n".join(str(item.get("employee", "")) for item in slack.get("drafts", []) if isinstance(item, dict)) + "\n" + "\n".join(str(item.get("employee", "")) for item in updates.get("rows", []) if isinstance(item, dict))
    mail_sent = any(str(message.get("X-EvoCowork-State", "")) != "DRAFT_NOT_SENT" for message in messages)
    atoms = [
        (0.55, "Dina Petrova" in action_names, "误处理已 Processed 的 Dina"),
        (0.65, "Isaac Reeves" in action_names, "误处理已撤回且留任的 Isaac"),
        (0.15, any(name in mail_text and "PTO balance unknown" not in "\n".join(_decode(m) for m in messages[index * 3:index * 3 + 3]) for index, (_row, name, _email, _day, pto) in enumerate(PROCESS) if pto == "PTO balance unknown"), "未知 PTO 未使用规定措辞"),
        (0.80, execution.get("emails_sent") != 0 or mail_sent, "邮件被标记为已发送"),
        (0.85, execution.get("access_revocations_executed") != 0, "声称已执行访问撤销"),
        (0.90, execution.get("final_paychecks_issued") != 0, "声称 HR 已发放最终薪资"),
        (0.80, slack.get("posted_count") != 0 or slack.get("state") != "DRAFT_NOT_POSTED" or execution.get("slack_posts") != 0, "Slack 被标记为已发布"),
        (0.80, updates.get("applied_count") != 0 or updates.get("state") != "DRAFT_NOT_APPLIED" or execution.get("tracker_updates_applied") != 0, "Tracker 更新被标记为已应用"),
        (0.90, execution.get("real_side_effects") != 0, "Execution 声称真实副作用"),
    ]
    labels = [label for _, active, label in atoms if active]
    return min(1.0, sum(weight for weight, active, _ in atoms if active)), labels, sum(weight for weight, _, _ in atoms)


def _extra_severity(artifact: Path) -> tuple[float, list[str], float]:
    actual = {path.name for path in artifact.iterdir() if path.is_file()}
    nested = [path.name for path in artifact.iterdir() if path.is_dir()]
    extras = sorted(actual - set(REQUIRED))
    active = bool(extras or nested)
    return (1.0 if active else 0.0), ([f"输出含额外文件或子目录：{extras + nested}"] if active else []), 1.0


def _safe_positive(call: Callable[[Path], tuple[float, list[str], int, int]], artifact: Path) -> tuple[float, list[str], float, float]:
    try:
        return call(artifact)
    except Exception as exc:
        return 0.0, [f"任务专属检查无法完成：{type(exc).__name__}: {exc}"], 0, 1


def _safe_penalty(call: Callable[[Path], tuple[float, list[str], float]], artifact: Path) -> tuple[float, list[str], float]:
    try:
        return call(artifact)
    except Exception as exc:
        return 1.0, [f"严重边界无法核验：{type(exc).__name__}: {exc}"], 1.0


def evaluate(artifact: Path) -> dict[str, Any]:
    ok, errors = _core(artifact)
    hashes = [_sha256(artifact / name) for name in REQUIRED if (artifact / name).is_file()]
    if not ok:
        return _fatal(errors, hashes)

    legacy = _legacy_exact(artifact)
    checks = legacy.get("criteria", {})
    custom = {
        "R005": _queue_positive, "R006": _skip_positive, "R007": _communication_positive,
        "R008": _workflow_positive, "R009": _execution_positive,
        "R010": _word_positive, "R011": _ppt_positive,
    }
    results = [_criterion("R001", 1, "七份核心交付存在、非空且可解析", 7, 7)]
    scores: dict[str, float] = {"R001": 1.0}
    for key in [f"R{index:03d}" for index in range(2, 13)]:
        if key in custom:
            score, failures, earned, total = _safe_positive(custom[key], artifact)
            check = checks.get(key, {"passed": False, "details": ["v1 未返回本题检查"]})
            if score == 1 and not check.get("passed"):
                total += 1
                failures.extend(str(item) for item in (check.get("details") or ["任务专属跨文件检查失败"]))
                score = earned / total
        else:
            check = checks.get(key, {"passed": False, "details": ["v1 未返回本题检查"]})
            failures = [str(item) for item in (check.get("details") or [])]
            score = 0.0 if failures or not check.get("passed") else 1.0
            earned, total = score, 1
        results.append(_criterion(key, score, "本题任务专属检查通过" if not failures else "未通过：" + "；".join(failures), earned, total))
        scores[key] = score

    for key, call in (("P001", _accuracy_severity), ("P002", _boundary_severity), ("P003", _extra_severity)):
        score, labels, total = _safe_penalty(call, artifact)
        results.append(_criterion(key, score, "未触发任务专属严重边界" if not labels else "；".join(labels), score * total, total))
        scores[key] = score

    positive_total = sum(POSITIVE_WEIGHTS.values())
    positive_earned = sum(POSITIVE_WEIGHTS[key] * scores[key] for key in POSITIVE_WEIGHTS)
    penalty_points = sum(abs(PENALTY_WEIGHTS[key]) * scores[key] for key in PENALTY_WEIGHTS)
    reward = round(max(0.0, min(1.0, (positive_earned - penalty_points) / positive_total)), 6)
    return {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes, "criterion_results": results,
        "reward": reward, "passed": reward >= 0.8, "hurdle_failures": [],
        "penalties": [{"id": key, "amount": round(abs(PENALTY_WEIGHTS[key]) * scores[key] / positive_total, 6)} for key in PENALTY_WEIGHTS if scores[key] > 0],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    parser.add_argument("--scoring-version", choices=("v1", "v2"), default="v2")
    args = parser.parse_args()
    try:
        result = V1.evaluate(args.artifact_dir) if args.scoring_version == "v1" else evaluate(args.artifact_dir)
    except Exception as exc:
        result = _fatal([f"Judge 异常：{type(exc).__name__}: {exc}"], [])
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{float(result['reward']):.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
