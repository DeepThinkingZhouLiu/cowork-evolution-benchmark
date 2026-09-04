#!/usr/bin/env python3
"""AutomationBench 5114 终止队列、PTO、撤回与零执行专属 Judge。"""
import argparse,json,mailbox,re
from pathlib import Path
from openpyxl import load_workbook
from office_checks import body_text,docx_text,package_safe,pptx_text,render,sha256
TASK="cowork-evo/automationbench-5114-termination-checklist-governance-package";HASHES={"HR_Ops_Slack.json":"2eec9fe8ac4d96cbf2c8c7b1f759ab706a34770f8125d747fdb45bf9623d8a36","Termination_Procedure.mbox":"9f3d3f5bd9d2fb90a1ec63f129cfc222bfda85cf8fce84980301e1f68bf38df9","Termination_Tracker.xlsx":"c5b11b596db2e8ff1274e5d27790a9ecb02796165b1a40ecbab83254720f6edc","Workflow_Context.json":"76c4a0657bc27f2fd830b10613ec7ceae3adcdc5798aeee0f822d27bf69bbc99"}
X="Termination_Checklist_Audit.xlsx";D="Termination_Checklist_Brief.docx";P="Termination_Checklist_Review.pptx";M="Termination_Communication_Drafts.mbox";S="HR_Ops_Termination_Slack_Drafts.json";U="Termination_Tracker_Update_Drafts.json";C="Termination_Execution_Control.json";REQUIRED={X,D,P,M,S,U,C};WEIGHTS={f"R{i:03d}":v for i,v in enumerate([6,8,10,10,10,8,9,8,8,8,7,8],1)};CAPS={k:(0 if k=="R001" else .25) for k in WEIGHTS if k!="R012"};ALL=["Aaron Wells","Bianca Morales","Chad Eriksson","Dina Petrova","Isaac Reeves","Ethan Park","Fiona Strand","George Tan","Holly Ng"];PROCESS=["Aaron Wells","Bianca Morales","Chad Eriksson","Ethan Park","Fiona Strand","George Tan","Holly Ng"]
class Checks:
 def __init__(self):self.v={k:{"passed":True,"details":[]} for k in WEIGHTS}
 def require(self,key,condition,detail):
  if not condition:self.v[key]["passed"]=False;self.v[key]["details"].append(detail) if detail not in self.v[key]["details"] else None
 def ok(self,key):return self.v[key]["passed"]
def root():return Path(__file__).resolve().parents[1]
def inp():return Path("/app/input_files") if Path("/app/input_files").is_dir() else root()/"data/input_files"
def rows(ws):return [tuple(c.value for c in row) for row in ws.iter_rows(min_row=2)]
def source(q):
 for name,digest in HASHES.items():q.require("R001",(inp()/name).is_file() and sha256(inp()/name)==digest,f"输入{name}")
 wb=load_workbook(inp()/"Termination_Tracker.xlsx",read_only=True,data_only=True);rs=rows(wb["Pending Terminations"]);q.require("R002",[r[1] for r in rs]==ALL and [r[0] for r in rs]==[2,3,4,5,10,6,7,8,9],"九行与源行号");q.require("R002",rs[3][6:8]==("Processed","Already completed") and rs[4][6]=="Pending" and "Rescinded" in rs[4][7] and [rs[i][5] for i in (1,5,8)]==[None,"N/A","TBD"],"已处理、撤回与PTO缺口")
 messages=list(mailbox.mbox(inp()/"Termination_Procedure.mbox",create=False));text=body_text(messages[0]) if len(messages)==1 else "";q.require("R002",all(token in text for token in ("it-ops@company.example.com","HR does NOT process final pay","payroll@company.example.com","PTO balance unknown","Update Status to 'Processed'","#hr-ops")),"终止程序")
def workbook(path,q):
 wb=load_workbook(path,data_only=False);expected=["Source_Tracker","Source_Messages","Policy_Rules","Evidence_Register","Checklist_Decisions","Process_Queue","Skip_Register","IT_Revocation_Drafts","Payroll_Request_Drafts","Tracker_Update_Drafts","Slack_Draft_Register","Control_Summary"];q.require("R003",wb.sheetnames==expected,"12表")
 formulas=[c.value for row in wb["Checklist_Decisions"].iter_rows(min_row=2,max_row=10) for c in row if isinstance(c.value,str) and c.value.startswith("=")];q.require("R004",len(formulas)==207,f"公式{len(formulas)}");joined="\n".join(formulas)
 for token in ('SEARCH("rescinded"','G2="Processed"','F2="N/A"','F2="TBD"','PTO balance unknown','SKIP_ALREADY_PROCESSED','SKIP_RESCINDED','PROCESS_TERMINATION_CHECKLIST'):q.require("R004",token in joined,f"公式缺{token}")
 process=rows(wb["Process_Queue"]);q.require("R005",[r[1] for r in process]==PROCESS and [r[3] for r in process]==["5 days","PTO balance unknown","12 days","PTO balance unknown","8 days","0 days","PTO balance unknown"],"七人队列与PTO")
 skips=rows(wb["Skip_Register"]);q.require("R006",[(r[1],r[4]) for r in skips]==[("Dina Petrova","SKIP_ALREADY_PROCESSED"),("Isaac Reeves","SKIP_RESCINDED")],"两跳过")
 q.require("R005",len(rows(wb["IT_Revocation_Drafts"]))==7 and len(rows(wb["Payroll_Request_Drafts"]))==7 and len(rows(wb["Tracker_Update_Drafts"]))==7,"三类七草稿")
def communications(artifact,q):
 messages=list(mailbox.mbox(artifact/M,create=False));q.require("R007",len(messages)==21,"邮件21")
 expected_employees=["aaron.wells@company.example.com","bianca.morales@company.example.com","chad.eriksson@company.example.com","ethan.park@company.example.com","fiona.strand@company.example.com","george.tan@company.example.com","holly.ng@company.example.com"]
 expected_days=["2026-04-18","2026-04-16","2026-04-20","2026-04-22","2026-04-25","2026-04-17","2026-04-30"];expected_pto=["5 days","PTO balance unknown","12 days","PTO balance unknown","8 days","0 days","PTO balance unknown"]
 for i,name in enumerate(PROCESS):
  triple=messages[i*3:i*3+3];q.require("R007",len(triple)==3 and [str(m.get("To")) for m in triple]==["it-ops@company.example.com","payroll@company.example.com",expected_employees[i]],f"{name}路由")
  texts=[body_text(m) for m in triple];q.require("R007",all(name in t and expected_days[i] in t and "DRAFT_NOT_SENT" in t for t in texts) and expected_pto[i] in texts[1] and "no access has been revoked" in texts[0] and "HR has not issued any payment" in texts[1],f"{name}日期、PTO与边界")
 # 员工收件人需要用来源精确核验，不依赖姓名推导。
 q.require("R007",[str(messages[i*3+2].get("To")) for i in range(7)]==expected_employees,"员工精确收件人")
 alltext="\n".join(body_text(m) for m in messages);q.require("R007","Dina Petrova" not in alltext and "Isaac Reeves" not in alltext,"跳过者误消息")
 slack=json.loads((artifact/S).read_text());q.require("R008",slack.get("channel")=="hr-ops" and slack.get("draft_count")==7 and slack.get("posted_count")==0 and slack.get("state")=="DRAFT_NOT_POSTED" and [x["employee"] for x in slack.get("drafts",[])]==PROCESS, "Slack七草稿")
 updates=json.loads((artifact/U).read_text());q.require("R008",updates.get("state")=="DRAFT_NOT_APPLIED" and updates.get("draft_count")==7 and updates.get("applied_count")==0 and [x["employee"] for x in updates.get("rows",[])]==PROCESS and [(x["employee"],x["decision"]) for x in updates.get("skipped",[])]==[("Dina Petrova","SKIP_ALREADY_PROCESSED"),("Isaac Reeves","SKIP_RESCINDED")],"台账草稿")
 control=json.loads((artifact/C).read_text());expected={"reviewed":9,"process_checklist":7,"already_processed_skipped":1,"rescinded_skipped":1,"it_drafts":7,"payroll_drafts":7,"employee_drafts":7,"email_drafts":21,"emails_sent":0,"slack_drafts":7,"slack_posts":0,"tracker_update_drafts":7,"tracker_updates_applied":0,"access_revocations_executed":0,"final_paychecks_issued":0,"real_side_effects":0};q.require("R009",all(control.get(k)==v for k,v in expected.items()),"执行边界")
def office(artifact,q):
 text,sections,tables=docx_text(artifact/D);q.require("R010",sections==3 and tables==3,"Word结构")
 for token in ("Seven pending terminations","Dina Petrova","Isaac Reeves","Bianca","Ethan","Holly","PTO balance unknown","does not issue final pay","No access was revoked","no source row was changed"):q.require("R010",token.lower() in text.lower(),f"Word缺{token}")
 q.require("R010","access was revoked, final pay was issued" not in text.lower(),"Word虚假执行")
 text,slides,tables,pictures=pptx_text(artifact/P);q.require("R011",slides==3 and tables==[1,1,1] and pictures==0,"PPT结构")
 for token in ("7 process","Dina","Isaac","Bianca","Ethan","Holly","DRAFT_NOT_SENT","DRAFT_NOT_POSTED","NOT_APPLIED"):q.require("R011",token in text,f"PPT缺{token}")
 for name in (X,D,P):safe,detail=package_safe(artifact/name);q.require("R012",safe,f"安全{name}:{detail}")
def finish(q,pages):
 hurdles=[k for k in CAPS if not q.ok(k)];base=sum(v for k,v in WEIGHTS.items() if q.ok(k))/100;cap=min([1]+[CAPS[k] for k in hurdles]);penalty=(.3 if any(not q.ok(k) for k in ("R002","R003","R004","R005","R006")) else 0)+(.3 if any(not q.ok(k) for k in ("R007","R008","R009")) else 0)+(.25 if any(not q.ok(k) for k in ("R001","R010","R011","R012")) else 0);reward=round(max(0,min(base,cap)-penalty),6);return {"task_id":TASK,"reward":reward,"passed":reward>=.8 and not hurdles,"criteria":q.v,"hurdle_failures":hurdles,"source_truth":{"reviewed":9,"process":7,"skip_processed":1,"skip_rescinded":1,"emails_sent":0,"revocations":0,"paychecks":0,"real_side_effects":0},"render_pages":pages}
def evaluate(artifact):
 q=Checks();source(q);files={p.name for p in artifact.iterdir() if p.is_file()} if artifact.is_dir() else set();q.require("R001",files==REQUIRED,f"集合{files}")
 if files!=REQUIRED:return finish(q,{})
 for fn,keys in ((lambda:workbook(artifact/X,q),("R003","R004","R005","R006","R012")),(lambda:communications(artifact,q),("R007","R008","R009")),(lambda:office(artifact,q),("R010","R011","R012"))):
  try:fn()
  except Exception as exc:
   for key in keys:q.require(key,False,str(exc))
 pages={}
 for name,expected in ((X,12),(D,3),(P,3)):
  try:count,text=render(artifact/name);pages[name]=count;q.require("R012",count==expected,f"页数{name}:{count}");q.require("R012",not re.search(r"#(?:VALUE|REF|NAME|DIV/0|N/A)",text,re.I),f"公式错{name}")
  except Exception as exc:q.require("R012",False,str(exc))
 return finish(q,pages)
def main():
 parser=argparse.ArgumentParser();parser.add_argument("--artifact-dir",type=Path,required=True);parser.add_argument("--result",type=Path,required=True);parser.add_argument("--reward-file",type=Path,required=True);args=parser.parse_args();result=evaluate(args.artifact_dir);args.result.parent.mkdir(parents=True,exist_ok=True);args.reward_file.parent.mkdir(parents=True,exist_ok=True);args.result.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n");args.reward_file.write_text(f"{result['reward']:.6f}\n");print(json.dumps(result,ensure_ascii=False,indent=2))
if __name__=="__main__":main()
