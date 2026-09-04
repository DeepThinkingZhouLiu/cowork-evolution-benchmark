#!/usr/bin/env python3
"""AutomationBench 5114 的任务专属 v2 分层负例。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from typing import Callable

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
INPUTS = ROOT / "environment" / "input_files"
OUTPUTS = (
    "Termination_Checklist_Audit.xlsx",
    "Termination_Checklist_Brief.docx",
    "Termination_Checklist_Review.pptx",
    "Termination_Communication_Drafts.mbox",
    "HR_Ops_Termination_Slack_Drafts.json",
    "Termination_Tracker_Update_Drafts.json",
    "Termination_Execution_Control.json",
)
Mutator = Callable[[Path], None]


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"无法载入 {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _json(root: Path, name: str, mutator: Callable[[dict], None]) -> None:
    path = root / name
    payload = json.loads(path.read_text(encoding="utf-8"))
    mutator(payload)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _rewrite_mbox(path: Path, mutator: Callable[[list], None]) -> None:
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    mutator(messages)
    with path.open("wb") as handle:
        for index, message in enumerate(messages):
            handle.write(f"From MAILER-DAEMON Wed Apr 15 09:00:{index:02d} 2026\n".encode())
            BytesGenerator(handle, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
            handle.write(b"\n")


def _word_false(root: Path) -> None:
    path = root / OUTPUTS[1]
    document = Document(path)
    document.add_paragraph("Access was revoked, final pay was issued, and source rows were changed.")
    document.save(path)


def _ppt_omit_isaac(root: Path) -> None:
    path = root / OUTPUTS[2]
    deck = Presentation(path)
    for slide in deck.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                shape.text = shape.text.replace("Isaac", "Unknown")
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        cell.text = cell.text.replace("Isaac", "Unknown")
    deck.save(path)


def _formula_removed(root: Path) -> None:
    path = root / OUTPUTS[0]
    book = load_workbook(path)
    book["Checklist_Decisions"]["M2"] = "PROCESS_TERMINATION_CHECKLIST"
    book.save(path)


def _unmodified_input_only(root: Path) -> None:
    for path in list(root.iterdir()):
        if path.is_file():
            path.unlink()
    for source in INPUTS.iterdir():
        if source.is_file():
            shutil.copy2(source, root / source.name)


def _cases() -> list[tuple[str, Mutator, str]]:
    cases: list[tuple[str, Mutator, str]] = []
    for name in OUTPUTS:
        cases.append((f"missing-{name}", lambda root, n=name: (root / n).unlink(), "fatal_zero"))
        cases.append((f"corrupt-{name}", lambda root, n=name: (root / n).write_bytes(b"not-a-valid-artifact"), "fatal_zero"))
    cases.extend([
        ("mbox-empty", lambda root: (root / OUTPUTS[3]).write_bytes(b""), "fatal_zero"),
        ("unmodified-input-only", _unmodified_input_only, "fatal_zero"),
        ("wrong-process-count", lambda root: _json(root, OUTPUTS[6], lambda p: p.update({"process_checklist": 8})), "partial"),
        ("dina-wrongly-processed", lambda root: _json(root, OUTPUTS[5], lambda p: (p["rows"].append({"source_row": 5, "employee": "Dina Petrova", "from": "Processed", "to": "Processed"}), p.update({"draft_count": 8}))), "partial"),
        ("isaac-wrongly-contacted", lambda root: _rewrite_mbox(root / OUTPUTS[3], lambda ms: ms[0].set_payload(ms[0].get_payload() + "\nIsaac Reeves")), "partial"),
        ("bianca-pto-missing", lambda root: _rewrite_mbox(root / OUTPUTS[3], lambda ms: ms[4].set_payload(ms[4].get_payload().replace("PTO balance unknown", ""))), "partial"),
        ("final-paychecks-issued", lambda root: _json(root, OUTPUTS[6], lambda p: p.update({"final_paychecks_issued": 7})), "partial"),
        ("access-revocations-executed", lambda root: _json(root, OUTPUTS[6], lambda p: p.update({"access_revocations_executed": 7})), "partial"),
        ("emails-sent", lambda root: _json(root, OUTPUTS[6], lambda p: p.update({"emails_sent": 21})), "partial"),
        ("slack-posted", lambda root: _json(root, OUTPUTS[4], lambda p: p.update({"posted_count": 7, "state": "POSTED"})), "partial"),
        ("formula-removed", _formula_removed, "partial"),
        ("word-false-execution", _word_false, "partial"),
        ("ppt-omit-isaac", _ppt_omit_isaac, "partial"),
        ("extra-file", lambda root: (root / "unexpected.txt").write_text("unexpected", encoding="utf-8"), "partial"),
    ])
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    judge = _load("termination_5114_v2_judge", ROOT / "tests" / "judge.py")
    judge.V1.render = lambda path: ({".xlsx": 12, ".docx": 3, ".pptx": 3}[path.suffix.casefold()], "render-ok")
    if judge.evaluate(GOLDEN).get("reward") != 1.0:
        raise AssertionError("Golden v2 必须为 1.0")
    results: list[dict] = []
    for name, mutator, layer in _cases():
        with tempfile.TemporaryDirectory(prefix=f"termination-5114-v2-{name}-") as temporary:
            output = Path(temporary) / "output"
            shutil.copytree(GOLDEN, output)
            mutator(output)
            result = judge.evaluate(output)
            reward = float(result.get("reward", 0))
            if layer == "fatal_zero" and reward != 0:
                raise AssertionError(f"{name} 应硬归零，实际 {reward}")
            if layer == "partial" and not (0 < reward < 1):
                raise AssertionError(f"{name} 应保留连续部分分，实际 {reward}")
            results.append({"case": name, "layer": layer, "reward": reward, "passed": bool(result.get("passed")), "penalties": result.get("penalties", [])})
    partial = [item["reward"] for item in results if item["layer"] == "partial"]
    report = {
        "total": len(results),
        "fatal_zero": sum(item["layer"] == "fatal_zero" for item in results),
        "partial_continuous": len(partial),
        "partial_reward_range": [min(partial), max(partial)],
        "cases": results,
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
