#!/usr/bin/env python3
"""只把 AutomationBench 5115 的公开初始状态静态化为本题输入。

本脚本只接受 ``hr.diversity_metrics_report``，逐项断言十条员工记录、
三条部门负责人记录、两封权威邮件、冻结时间、Slack 频道和九条上游断言。
它不是通用制题器，不会创造员工、部门、完成状态、负责人或隐私规则。
"""

from __future__ import annotations

import hashlib
import importlib.util
import json
import shutil
import sys
import types
import zipfile
from datetime import datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
UPSTREAM = Path(__file__).with_name("upstream-hr-tasks.py")
SNAPSHOT = Path(__file__).with_name("upstream-task-5115.json")
LICENSE = Path(__file__).with_name("LICENSE-AutomationBench-MIT")
DATA_DIR = ROOT / "data" / "input_files"
ENV_DIR = ROOT / "environment" / "input_files"
FIXED_ZIP_TIME = (2026, 4, 15, 9, 0, 0)
SOURCE_SHA256 = "c173a36a03f43096bd122577584aa32f9df379d8760aa46c43be04c85d101779"
LICENSE_SHA256 = "c6f9d87aff6d627d5cd78cb09364dafed7d6f1c4bb7a91183d57377157e96ba4"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def freeze_upstream_assets() -> None:
    source_task = ROOT.parent / "automationbench-5113-i9-compliance-control-package" / "adaptation"
    if not UPSTREAM.is_file():
        source = source_task / "upstream-hr-tasks.py"
        if not source.is_file() or sha256(source) != SOURCE_SHA256:
            raise FileNotFoundError("找不到哈希匹配的 AutomationBench HR 冻结源文件")
        shutil.copy2(source, UPSTREAM)
    if sha256(UPSTREAM) != SOURCE_SHA256:
        raise ValueError("AutomationBench HR 冻结源文件哈希错误")
    if not LICENSE.is_file():
        source = source_task / "LICENSE-AutomationBench-MIT"
        if not source.is_file() or sha256(source) != LICENSE_SHA256:
            raise FileNotFoundError("找不到哈希匹配的 AutomationBench LICENSE")
        shutil.copy2(source, LICENSE)
    if sha256(LICENSE) != LICENSE_SHA256:
        raise ValueError("AutomationBench LICENSE 哈希错误")


def load_task() -> dict:
    datasets = types.ModuleType("datasets")
    datasets.Dataset = object
    sys.modules["datasets"] = datasets
    for name in ("automationbench", "automationbench.domains", "automationbench.domains.hr"):
        sys.modules[name] = types.ModuleType(name)
    noise = types.ModuleType("automationbench.domains.hr._noise")
    noise.apply_noise = lambda value: value
    sys.modules[noise.__name__] = noise
    spec = importlib.util.spec_from_file_location("automationbench_hr_tasks_frozen_5115", UPSTREAM)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入冻结的 AutomationBench HR 源文件")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    task = module.get_task_5115()
    if task.get("example_id") != 5115 or task.get("task") != "hr.diversity_metrics_report":
        raise ValueError("冻结函数不再是 AutomationBench 5115")
    return task


def assert_shape(task: dict) -> None:
    prompt = task["prompt"][1]["content"]
    assert all(token in prompt for token in ("workforce diversity summary", "each department", "below 70%", "specific employees", "Slack summary"))
    state = task["info"]["initial_state"]
    assert state["meta"] == {"schema_version": "0.1.0", "current_time": "2026-04-15T09:00:00Z"}
    messages = state["gmail"]["messages"]
    assert [message["id"] for message in messages] == ["msg_selfid_privacy_5115", "msg_dept_correction_5115"]
    assert all(token in messages[0]["body_plain"] for token in ("aggregate percentages ONLY", "Under NO circumstances", "must be refused"))
    assert all(token in messages[1]["body_plain"] for token in ("Ingrid Larsson", "Sales to Engineering", "Jake Morrison", "Support to Sales", "current departments"))
    spreadsheet = state["google_sheets"]["spreadsheets"][0]
    assert spreadsheet["id"] == "ss_diversity_5115" and spreadsheet["title"] == "Workforce Demographics"
    assert [worksheet["title"] for worksheet in spreadsheet["worksheets"]] == ["Self-ID Status", "Department Heads"]
    employees = spreadsheet["worksheets"][0]["rows"]
    assert [row["row_id"] for row in employees] == list(range(2, 12))
    assert [row["cells"]["Employee"] for row in employees] == ["Ada Chen", "Ben Osei", "Clara Dubois", "Derek Huang", "Elena Rossi", "Femi Adeyemi", "Greta Nilsson", "Ingrid Larsson", "Jake Morrison", "Karen Blake"]
    assert [row["cells"]["Self-ID Completed"] for row in employees] == ["Yes", "Yes", "No", "Yes", "No", "Yes", "No", "Yes", "No", "Yes"]
    heads = spreadsheet["worksheets"][1]["rows"]
    assert [(row["cells"]["Department"], row["cells"]["Email"]) for row in heads] == [("Engineering", "vp-eng@company.example.com"), ("Sales", "vp-sales@company.example.com"), ("Support", "support-director@company.example.com")]
    assert state["slack"] == {"channels": [{"id": "C_HRL_5115", "name": "hr-leadership", "is_private": False}], "messages": []}
    assert len(task["info"]["assertions"]) == 9
    assertion_text = json.dumps(task["info"]["assertions"])
    assert "Clara Dubois" in assertion_text and "Greta Nilsson" in assertion_text
    assert "Elena Rossi" not in assertion_text and "Jake Morrison" not in assertion_text


def normalize_xlsx(path: Path) -> None:
    temporary = path.with_suffix(".normalized.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    temporary.replace(path)


def style_sheet(sheet, widths: dict[str, float]) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = sheet.dimensions


def add_table(sheet, name: str) -> None:
    table = Table(displayName=name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)


def write_workforce(state: dict, target: Path) -> None:
    source = state["google_sheets"]["spreadsheets"][0]
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.creator = "AutomationBench 5115 public fixture"
    workbook.title = source["title"]
    workbook.subject = "Static Office representation of workforce self-ID status and department heads"
    workbook.description = "No employee, department, completion state, head or email was added beyond AutomationBench 5115."
    employees = workbook.create_sheet("Self-ID Status")
    employees.append(["Source Row", "Employee", "Department", "Self-ID Completed", "Notes"])
    for item in source["worksheets"][0]["rows"]:
        cells = item["cells"]
        employees.append([item["row_id"], cells["Employee"], cells["Department"], cells["Self-ID Completed"], cells["Notes"]])
    add_table(employees, "SelfIDStatusSourceInput")
    style_sheet(employees, {"A": 12, "B": 27, "C": 24, "D": 22, "E": 28})
    heads = workbook.create_sheet("Department Heads")
    heads.append(["Source Row", "Department", "Head", "Email"])
    for item in source["worksheets"][1]["rows"]:
        cells = item["cells"]
        heads.append([item["row_id"], cells["Department"], cells["Head"], cells["Email"]])
    add_table(heads, "DepartmentHeadsSourceInput")
    style_sheet(heads, {"A": 12, "B": 24, "C": 28, "D": 38})
    workbook.save(target)
    normalize_xlsx(target)


def message_bytes(record: dict) -> bytes:
    message = EmailMessage(policy=SMTP)
    message["From"] = record["from_"]
    message["To"] = ", ".join(record["to"])
    message["Subject"] = record["subject"]
    parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
    message["Date"] = parsed.strftime("%a, %d %b %Y %H:%M:%S +0000")
    message["Message-ID"] = f"<{record['id']}@automationbench.invalid>"
    message["X-AutomationBench-Message-ID"] = record["id"]
    message["X-AutomationBench-Thread-ID"] = record["thread_id"]
    message["X-AutomationBench-Labels"] = ",".join(record.get("label_ids", []))
    message.set_content(record["body_plain"], charset="utf-8")
    buffer = BytesIO()
    BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
    return buffer.getvalue()


def write_mbox(messages: list[dict], target: Path) -> None:
    chunks: list[bytes] = []
    for record in messages:
        parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
        chunks.append(f"From {record['from_']} {parsed.strftime('%a %b %d %H:%M:%S %Y')}\n".encode("ascii"))
        chunks.append(message_bytes(record).replace(b"\r\n", b"\n"))
        chunks.append(b"\n")
    target.write_bytes(b"".join(chunks))


def main() -> None:
    freeze_upstream_assets()
    task = load_task()
    assert_shape(task)
    state = task["info"]["initial_state"]
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    ENV_DIR.mkdir(parents=True, exist_ok=True)
    SNAPSHOT.write_text(json.dumps(task, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_workforce(state, DATA_DIR / "Workforce_Demographics.xlsx")
    write_mbox(state["gmail"]["messages"], DATA_DIR / "Diversity_Context.mbox")
    (DATA_DIR / "Slack_Channels.json").write_text(json.dumps(state["slack"], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (DATA_DIR / "Workflow_Context.json").write_text(json.dumps({"meta": state["meta"]}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    for source in sorted(DATA_DIR.iterdir()):
        if source.is_file():
            shutil.copy2(source, ENV_DIR / source.name)


if __name__ == "__main__":
    main()
