#!/usr/bin/env python3
"""只静态化 AutomationBench 5116；不是通用造题器。"""
import hashlib,importlib.util,json,shutil,sys,types
from datetime import date,datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.worksheet.table import Table,TableStyleInfo
ROOT=Path(__file__).resolve().parents[1];UP=Path(__file__).with_name("upstream-hr-tasks.py");SNAP=Path(__file__).with_name("upstream-task-5116.json");DATA=ROOT/"data/input_files";ENV=ROOT/"environment/input_files";UPSTREAM_SHA="c173a36a03f43096bd122577584aa32f9df379d8760aa46c43be04c85d101779"
def load_task():
 if hashlib.sha256(UP.read_bytes()).hexdigest()!=UPSTREAM_SHA:raise ValueError("上游哈希错误")
 d=types.ModuleType("datasets");d.Dataset=object;sys.modules["datasets"]=d
 for name in ("automationbench","automationbench.domains","automationbench.domains.hr"):sys.modules[name]=types.ModuleType(name)
 noise=types.ModuleType("automationbench.domains.hr._noise");noise.apply_noise=lambda x:x;sys.modules[noise.__name__]=noise
 spec=importlib.util.spec_from_file_location("automationbench_5116",UP);module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module);task=module.get_task_5116();assert task["example_id"]==5116 and task["task"]=="hr.buddy_assignment";return task
def audit(task):
 info=task["info"];state=info["initial_state"];assert state["meta"]["current_time"]=="2026-04-15T09:00:00Z" and len(info["assertions"])==7
 sheets=state["google_sheets"]["spreadsheets"][0]["worksheets"];assert [s["title"] for s in sheets]==["New Hires","Buddy Pool","Department Heads"] and len(sheets[0]["rows"])==3 and len(sheets[1]["rows"])==7 and len(sheets[2]["rows"])==3
 assert [r["cells"]["New Hire"] for r in sheets[0]["rows"]]==["Li Wei","Li Wei-Chen","Nora Lindgren"] and [r["cells"]["Name"] for r in sheets[1]["rows"]]==["Carlos Mendoza","Fatima Al-Rashid","James Whitfield","Sandra Liu","Olga Petrova","Brandon Osei","Ava Chen"]
 policy=state["gmail"]["messages"][0]["body_plain"];assert "SAME department" in policy and "before April 2025" in policy and "longest tenure" in policy and "Do NOT assign a cross-department buddy" in policy
def write_mbox(messages,path):
 chunks=[]
 for record in messages:
  msg=EmailMessage(policy=SMTP);msg["From"]=record["from_"];msg["To"]=", ".join(record["to"]);msg["Subject"]=record["subject"];dt=datetime.fromisoformat(record["date"].replace("Z","+00:00"));msg["Date"]=dt.strftime("%a, %d %b %Y %H:%M:%S +0000");msg["Message-ID"]=f"<{record['id']}@automationbench.invalid>";msg["X-AutomationBench-Message-ID"]=record["id"];msg.set_content(record["body_plain"],charset="utf-8");buf=BytesIO();BytesGenerator(buf,policy=SMTP,mangle_from_=True,maxheaderlen=0).flatten(msg);chunks.extend([f"From {record['from_']} {dt.strftime('%a %b %d %H:%M:%S %Y')}\n".encode(),buf.getvalue().replace(b"\r\n",b"\n"),b"\n"])
 path.write_bytes(b"".join(chunks))
def styled_sheet(ws,widths,table_name):
 for i,w in enumerate(widths,1):ws.column_dimensions[ws.cell(1,i).column_letter].width=w
 ws.freeze_panes="A2";ws.sheet_view.showGridLines=False
 for cell in ws[1]:cell.fill=PatternFill("solid",fgColor="17365D");cell.font=Font(name="Carlito",bold=True,color="FFFFFF");cell.alignment=Alignment(horizontal="center",wrap_text=True)
 for row in ws.iter_rows(min_row=2):
  for cell in row:cell.font=Font(name="Carlito",size=9);cell.alignment=Alignment(vertical="top",wrap_text=True)
 table=Table(displayName=table_name,ref=ws.dimensions);table.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True);ws.add_table(table);ws.page_setup.orientation="landscape";ws.page_setup.fitToWidth=1;ws.sheet_properties.pageSetUpPr.fitToPage=True
def write_xlsx(state,path):
 source=state["google_sheets"]["spreadsheets"][0]["worksheets"];wb=Workbook();wb.remove(wb.active)
 for index,(title,headers,widths,table) in enumerate((("New Hires",["Source Row","New Hire","New Hire Email","Department","Start Date"],[10,24,38,22,18],"NewHiresSource"),("Buddy Pool",["Source Row","Name","Email","Department","Hire Date","Currently Assigned","Notes"],[10,24,38,22,18,22,58],"BuddyPoolSource"),("Department Heads",["Source Row","Department","Head Email"],[10,22,38],"DepartmentHeadsSource"))):
  ws=wb.create_sheet(title);ws.append(headers)
  for row in source[index]["rows"]:
   c=row["cells"]
   if index==0:values=[row["row_id"],c["New Hire"],c["New Hire Email"],c["Department"],date.fromisoformat(c["Start Date"])]
   elif index==1:values=[row["row_id"],c["Name"],c["Email"],c["Department"],date.fromisoformat(c["Hire Date"]),c["Currently Assigned"],c["Notes"]]
   else:values=[row["row_id"],c["Department"],c["Head Email"]]
   ws.append(values)
  styled_sheet(ws,widths,table)
  if index<2:
   for row in ws.iter_rows(min_row=2):row[4].number_format="yyyy-mm-dd"
 wb.save(path)
def main():
 task=load_task();audit(task);state=task["info"]["initial_state"];DATA.mkdir(parents=True,exist_ok=True);ENV.mkdir(parents=True,exist_ok=True);snapshot=json.dumps(task,ensure_ascii=False,indent=2)+"\n"
 if SNAP.exists() and SNAP.read_text()!=snapshot:raise ValueError("单题快照漂移")
 SNAP.write_text(snapshot);write_xlsx(state,DATA/"Onboarding_Buddy_Tracker.xlsx");write_mbox(state["gmail"]["messages"],DATA/"Buddy_Program_Policy.mbox");(DATA/"Buddy_Slack_Users.json").write_text(json.dumps(state["slack"],ensure_ascii=False,indent=2)+"\n");(DATA/"Workflow_Context.json").write_text(json.dumps({"meta":state["meta"]},indent=2)+"\n")
 for item in DATA.iterdir():
  if item.is_file():shutil.copy2(item,ENV/item.name)
if __name__=="__main__":main()
