#!/usr/bin/env python3
"""AutomationBench 5116 的 v2 连续评分入口；原严格口径保存在 judge_v1.py。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import mailbox
import sys
import tempfile
from email.utils import getaddresses
from pathlib import Path
from typing import Any, Callable

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


TASK_ID = "cowork-evo/automationbench-5116-buddy-assignment-governance-package"
RUBRIC_VERSION = "automationbench-5116-buddy-v0.2"
JUDGE_VERSION = "automationbench-5116-criterion-partial-v2.0"
XLSX = "Buddy_Assignment_Audit.xlsx"
DOCX = "Buddy_Assignment_Brief.docx"
PPTX = "Buddy_Assignment_Review.pptx"
NEW_HIRE_MBOX = "New_Hire_Buddy_Drafts.mbox"
DM_JSON = "Buddy_Slack_DM_Drafts.json"
HEAD_MBOX = "Department_Head_Escalation_Draft.mbox"
EXECUTION_JSON = "Buddy_Execution_Control.json"
REQUIRED = (XLSX, DOCX, PPTX, NEW_HIRE_MBOX, DM_JSON, HEAD_MBOX, EXECUTION_JSON)
POSITIVE_WEIGHTS = {
    "R001": 5, "R002": 8, "R003": 9, "R004": 11,
    "R005": 12, "R006": 9, "R007": 8, "R008": 9,
    "R009": 8, "R010": 6, "R011": 6, "R012": 9,
}
PENALTY_WEIGHTS = {"P001": -35, "P002": -35, "P003": -10}
HIRES = [
    (2, "Li Wei", "li.wei@company.example.com", "Engineering", "2026-04-21"),
    (3, "Li Wei-Chen", "li.weichen@company.example.com", "Marketing", "2026-04-21"),
    (4, "Nora Lindgren", "nora.lindgren@company.example.com", "Sales", "2026-04-21"),
]
POOL = [
    (2, "Carlos Mendoza", "carlos.mendoza@company.example.com", "Engineering", "2023-03-20", "No", ""),
    (3, "Fatima Al-Rashid", "fatima.alrashid@company.example.com", "Engineering", "2024-06-15", "No", "Out until May 15 - parental leave"),
    (4, "James Whitfield", "james.whitfield@company.example.com", "Engineering", "2025-08-01", "No", ""),
    (5, "Sandra Liu", "sandra.liu@company.example.com", "Marketing", "2022-09-10", "No", ""),
    (6, "Olga Petrova", "olga.petrova@company.example.com", "Marketing", "2025-06-20", "No", ""),
    (7, "Brandon Osei", "brandon.osei@company.example.com", "Sales", "2025-11-01", "No", ""),
    (8, "Ava Chen", "ava.chen@company.example.com", "Sales", "2026-01-15", "No", ""),
]
HEADS = [
    (2, "Engineering", "vp-eng@company.example.com"),
    (3, "Marketing", "vp-mkt@company.example.com"),
    (4, "Sales", "vp-sales@company.example.com"),
]
PAIRS = [
    ("Li Wei", "Engineering", "Carlos Mendoza", "carlos.mendoza@company.example.com", "U_CARLOS", "2026-04-21", "DRAFT_NOT_APPLIED"),
    ("Li Wei-Chen", "Marketing", "Sandra Liu", "sandra.liu@company.example.com", "U_SANDRA", "2026-04-21", "DRAFT_NOT_APPLIED"),
]
EXECUTION_EXPECTED = {
    "task_id": "automationbench-5116",
    "new_hires_reviewed": 3,
    "assigned_pairs": 2,
    "head_escalations": 1,
    "new_hire_email_drafts": 2,
    "emails_sent": 0,
    "buddy_dm_drafts": 2,
    "buddy_dms_delivered": 0,
    "cross_department_assignments": 0,
    "assignments_applied": 0,
    "real_side_effects": 0,
}
DM_EXPECTED = {
    "task_id": "automationbench-5116",
    "draft_count": 2,
    "delivered_count": 0,
    "state": "DRAFT_NOT_DELIVERED",
    "drafts": [
        {
            "recipient_id": "U_CARLOS",
            "buddy": "Carlos Mendoza",
            "new_hire": "Li Wei",
            "text": "You are the proposed onboarding buddy for Li Wei in Engineering, starting 2026-04-21. DRAFT_NOT_DELIVERED.",
        },
        {
            "recipient_id": "U_SANDRA",
            "buddy": "Sandra Liu",
            "new_hire": "Li Wei-Chen",
            "text": "You are the proposed onboarding buddy for Li Wei-Chen in Marketing, starting 2026-04-21. DRAFT_NOT_DELIVERED.",
        },
    ],
}


def _load_v1():
    path = Path(__file__).with_name("judge_v1.py")
    spec = importlib.util.spec_from_file_location("automationbench_5116_judge_v1", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v1 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


V1 = _load_v1()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _norm(value: Any) -> Any:
    if value is None:
        return ""
    if hasattr(value, "date") and not isinstance(value, str):
        return value.date().isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _rows(sheet, start: int = 2, end: int | None = None, columns: int | None = None) -> list[tuple[Any, ...]]:
    end = end or sheet.max_row
    columns = columns or sheet.max_column
    return [
        tuple(_norm(sheet.cell(row, column).value) for column in range(1, columns + 1))
        for row in range(start, end + 1)
    ]


def _messages(path: Path) -> list:
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    return messages


def _body(message) -> str:
    payload = message.get_payload(decode=True)
    if payload is None:
        return str(message.get_payload()).strip()
    return payload.decode(message.get_content_charset() or "utf-8", errors="replace").strip()


def _addresses(message, header: str) -> set[str]:
    return {address.casefold() for _, address in getaddresses(message.get_all(header, [])) if address}


def _formula(value: Any, *tokens: str) -> bool:
    text = str(value or "")
    return text.startswith("=") and all(token.casefold() in text.casefold() for token in tokens)


def _core(artifact: Path) -> tuple[bool, list[str]]:
    errors: list[str] = []
    for name in REQUIRED:
        path = artifact / name
        if not path.is_file() or path.stat().st_size == 0:
            errors.append(f"核心交付缺失或为空：{name}")
            continue
        try:
            if name == XLSX:
                load_workbook(path, read_only=True, data_only=False).close()
            elif name == DOCX:
                Document(path)
            elif name == PPTX:
                if not Presentation(path).slides:
                    raise ValueError("PPT 无页面")
            elif name in (NEW_HIRE_MBOX, HEAD_MBOX):
                messages = _messages(path)
                if not messages or not all(message.get("To") and message.get("Subject") for message in messages):
                    raise ValueError("MBOX 无有效草稿")
            else:
                payload = json.loads(path.read_text(encoding="utf-8"))
                if not isinstance(payload, dict):
                    raise ValueError("JSON 顶层不是对象")
        except Exception as exc:
            errors.append(f"核心交付无法解析：{name}: {exc}")
    return not errors, errors


def _criterion(key: str, score: float, evidence: str, earned: float = 0, total: float = 1) -> dict[str, Any]:
    return {
        "criterion_id": key,
        "score": round(max(0.0, min(1.0, score)), 6),
        "evidence": evidence,
        "method": "deterministic",
        "raw": {"earned_units": round(earned, 6), "total_units": round(total, 6)},
    }


def _fatal(errors: list[str], hashes: list[str]) -> dict[str, Any]:
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes or ["0" * 64],
        "criterion_results": [_criterion(key, 0, "；".join(errors), 0, 0) for key in (*POSITIVE_WEIGHTS, *PENALTY_WEIGHTS)],
        "reward": 0.0,
        "passed": False,
        "hurdle_failures": ["R001"],
        "penalties": [],
    }


def _ratio(findings: list[tuple[bool, str]]) -> tuple[float, list[str], int, int]:
    earned = sum(bool(ok) for ok, _ in findings)
    return earned / len(findings), [label for ok, label in findings if not ok], earned, len(findings)


def _leaf_findings(observed: Any, expected: Any, path: str) -> list[tuple[bool, str]]:
    if isinstance(expected, dict):
        findings = [(isinstance(observed, dict) and set(observed) == set(expected), f"{path} 字段集合错误")]
        source = observed if isinstance(observed, dict) else {}
        for key, value in expected.items():
            findings.extend(_leaf_findings(source.get(key), value, f"{path}.{key}"))
        return findings
    if isinstance(expected, list):
        findings = [(isinstance(observed, list) and len(observed) == len(expected), f"{path} 列表长度错误")]
        source = observed if isinstance(observed, list) else []
        for index, value in enumerate(expected):
            findings.extend(_leaf_findings(source[index] if index < len(source) else None, value, f"{path}[{index}]"))
        return findings
    return [(observed == expected, f"{path} 应为 {expected!r}")]


def _legacy_exact(artifact: Path) -> dict[str, Any]:
    with tempfile.TemporaryDirectory(prefix="buddy-5116-v1-") as temporary:
        exact = Path(temporary)
        for name in REQUIRED:
            (exact / name).symlink_to((artifact / name).resolve())
        return V1.evaluate(exact)


def _source_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    book = load_workbook(artifact / XLSX, data_only=True)
    observed_sets = [
        (_rows(book["Source_New_Hires"], 2, 4, 5), HIRES, "Source_New_Hires"),
        (_rows(book["Source_Buddy_Pool"], 2, 8, 7), POOL, "Source_Buddy_Pool"),
        (_rows(book["Source_Department_Heads"], 2, 4, 3), HEADS, "Source_Department_Heads"),
    ]
    policy_rows = _rows(book["Source_Policy"], 2, 2, 2)
    evidence = _rows(book["Evidence_Register"], 2, 5, 3)
    book.close()
    findings: list[tuple[bool, str]] = []
    for observed_rows, expected_rows, label in observed_sets:
        for row_index, expected in enumerate(expected_rows):
            observed = observed_rows[row_index] if row_index < len(observed_rows) else ()
            for column, value in enumerate(expected):
                findings.append((observed[column] == value if column < len(observed) else False, f"{label} 第{row_index + 2}行第{column + 1}列错误"))
    source_policy = _messages(V1.inp() / "Buddy_Program_Policy.mbox")
    body = _body(source_policy[0]) if len(source_policy) == 1 else ""
    findings.extend([
        (len(policy_rows) == 1 and policy_rows[0][0] == "msg_buddy_policy_5116" and str(policy_rows[0][1]).strip() == body.strip(), "Source_Policy 未保留权威政策全文"),
        (len(evidence) == 4, "Evidence_Register 不是四个冻结输入"),
    ])
    for filename, digest in V1.HASHES.items():
        findings.append((any(row == (filename, "Frozen 5116 source", digest) for row in evidence), f"{filename} 证据哈希错误"))
    for token in ("SAME department", "before April 2025", "ONE new hire", "longest tenure", "Do NOT assign a cross-department buddy"):
        findings.append((token in body, f"Buddy政策缺少 {token}"))
    return _ratio(findings)


def _formula_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    book = load_workbook(artifact / XLSX, data_only=False)
    expected_sheets = [
        "Source_New_Hires", "Source_Buddy_Pool", "Source_Department_Heads", "Source_Policy",
        "Evidence_Register", "Buddy_Eligibility", "Assignment_Decisions", "Assigned_Pairs",
        "Unfilled_Assignments", "New_Hire_Email_Drafts", "Buddy_DM_Drafts", "Control_Summary",
    ]
    findings: list[tuple[bool, str]] = [(book.sheetnames == expected_sheets, "Excel十二表名称或顺序错误")]
    eligibility = book["Buddy_Eligibility"]
    for row in range(2, 9):
        for column in range(1, 8):
            letter = chr(64 + column)
            findings.append((_formula(eligibility.cell(row, column).value, f"Source_Buddy_Pool!{letter}{row}"), f"Buddy_Eligibility!{letter}{row} 未链接来源"))
        dependencies = {
            8: (f"E{row}", "DATE(2025,4,1)"),
            9: (f"F{row}", "No"),
            10: ("SEARCH", "Out until", f"G{row}"),
            11: ("AND", f"H{row}", f"I{row}", f"J{row}"),
            12: (f"K{row}", "ELIGIBLE", "ON_LEAVE", "TENURE_UNDER_ONE_YEAR"),
            13: (f"B{row}", f"D{row}", f"E{row}", f"K{row}"),
        }
        for column, tokens in dependencies.items():
            findings.append((_formula(eligibility.cell(row, column).value, *tokens), f"Buddy_Eligibility 第{row}行第{column}列公式错误"))
    decisions = book["Assignment_Decisions"]
    for row in range(2, 5):
        for column in range(1, 6):
            letter = chr(64 + column)
            findings.append((_formula(decisions.cell(row, column).value, f"Source_New_Hires!{letter}{row}"), f"Assignment_Decisions!{letter}{row} 未链接新人来源"))
        dependencies = {
            6: (f"D{row}", "Engineering", "Marketing", "Buddy_Eligibility!B2", "Buddy_Eligibility!B5"),
            7: (f"F{row}", "Buddy_Eligibility!C2", "Buddy_Eligibility!C5"),
            8: (f"F{row}", "Carlos Mendoza", "U_CARLOS", "Sandra Liu", "U_SANDRA"),
            9: (f"D{row}", "Source_Department_Heads!C2", "Source_Department_Heads!C3", "Source_Department_Heads!C4"),
            10: (f"F{row}", "ESCALATE_DEPARTMENT_HEAD", "ASSIGN_BUDDY"),
            11: (f"J{row}", "DRAFT_NOT_SENT", "NOT_REQUIRED"),
            12: (f"J{row}", "DRAFT_NOT_DELIVERED", "NOT_REQUIRED"),
            13: (f"J{row}", "DRAFT_NOT_SENT", "NOT_REQUIRED"),
            14: ("0",), 15: ("0",), 16: ("0",),
            17: (f"B{row}", f"D{row}", f"J{row}", f"F{row}"),
        }
        for column, tokens in dependencies.items():
            findings.append((_formula(decisions.cell(row, column).value, *tokens), f"Assignment_Decisions 第{row}行第{column}列公式错误"))
    book.close()
    return _ratio(findings)


def _eligibility_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    book = load_workbook(artifact / XLSX, data_only=True)
    pool = _rows(book["Source_Buddy_Pool"], 2, 8, 7)
    book.close()
    expected_reasons = {
        "Carlos Mendoza": "ELIGIBLE", "Fatima Al-Rashid": "ON_LEAVE",
        "James Whitfield": "TENURE_UNDER_ONE_YEAR", "Sandra Liu": "ELIGIBLE",
        "Olga Petrova": "TENURE_UNDER_ONE_YEAR", "Brandon Osei": "TENURE_UNDER_ONE_YEAR",
        "Ava Chen": "TENURE_UNDER_ONE_YEAR",
    }
    doc = Document(artifact / DOCX)
    rows = [tuple(cell.text for cell in row.cells) for table in doc.tables for row in table.rows]
    findings: list[tuple[bool, str]] = []
    for index, expected in enumerate(POOL):
        observed = pool[index] if index < len(pool) else ()
        for column, value in enumerate(expected):
            findings.append((observed[column] == value if column < len(observed) else False, f"{expected[1]} 来源字段{column + 1}错误"))
        findings.append(((expected[1], expected[3], expected[4], expected_reasons[expected[1]]) in rows, f"{expected[1]} 资格结论错误"))
    return _ratio(findings)


def _assignment_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    book = load_workbook(artifact / XLSX, data_only=True)
    pairs = _rows(book["Assigned_Pairs"], 2, 3, 7)
    unfilled = _rows(book["Unfilled_Assignments"], 2, 2, 5)
    emails = _rows(book["New_Hire_Email_Drafts"], 2, 3, 4)
    dms = _rows(book["Buddy_DM_Drafts"], 2, 3, 4)
    book.close()
    expected_unfilled = ("Nora Lindgren", "Sales", "vp-sales@company.example.com", "No same-department buddy meets tenure and availability rules", "ESCALATION_DRAFT")
    expected_emails = [
        ("Li Wei", "li.wei@company.example.com", "Carlos Mendoza", "DRAFT_NOT_SENT"),
        ("Li Wei-Chen", "li.weichen@company.example.com", "Sandra Liu", "DRAFT_NOT_SENT"),
    ]
    expected_dms = [
        ("Carlos Mendoza", "U_CARLOS", "Li Wei", "DRAFT_NOT_DELIVERED"),
        ("Sandra Liu", "U_SANDRA", "Li Wei-Chen", "DRAFT_NOT_DELIVERED"),
    ]
    findings: list[tuple[bool, str]] = []
    for label, observed_rows, expected_rows in (
        ("Assigned_Pairs", pairs, PAIRS),
        ("New_Hire_Email_Drafts", emails, expected_emails),
        ("Buddy_DM_Drafts", dms, expected_dms),
    ):
        findings.append((len(observed_rows) == len(expected_rows), f"{label} 行数错误"))
        for index, expected in enumerate(expected_rows):
            observed = observed_rows[index] if index < len(observed_rows) else ()
            for column, value in enumerate(expected):
                findings.append((observed[column] == value if column < len(observed) else False, f"{label} 第{index + 2}行第{column + 1}列错误"))
    for column, value in enumerate(expected_unfilled):
        observed = unfilled[0] if unfilled else ()
        findings.append((observed[column] == value if column < len(observed) else False, f"Nora缺口第{column + 1}字段错误"))
    findings.append((len(unfilled) == 1, "缺口登记不是恰好Nora一人"))
    return _ratio(findings)


def _new_hire_email_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    messages = _messages(artifact / NEW_HIRE_MBOX)
    expected = [
        ("li.wei@company.example.com", "Your onboarding buddy: Carlos Mendoza", "<buddy-newhire-5116-01@draft.invalid>", "Hello Li Wei,\n\nYour onboarding buddy is Carlos Mendoza from Engineering. Your start date is 2026-04-21. DRAFT_NOT_SENT."),
        ("li.weichen@company.example.com", "Your onboarding buddy: Sandra Liu", "<buddy-newhire-5116-02@draft.invalid>", "Hello Li Wei-Chen,\n\nYour onboarding buddy is Sandra Liu from Marketing. Your start date is 2026-04-21. DRAFT_NOT_SENT."),
    ]
    findings: list[tuple[bool, str]] = [(len(messages) == 2, "新人MBOX不是恰好两封")]
    for index, wanted in enumerate(expected):
        message = messages[index] if index < len(messages) else None
        if message is None:
            findings.extend((False, f"缺少第{index + 1}封新人草稿字段{part}") for part in range(8))
            continue
        to, subject, message_id, body = wanted
        findings.extend([
            (_addresses(message, "From") == {"hr-ops@company.example.com"}, f"第{index + 1}封新人草稿From错误"),
            (_addresses(message, "To") == {to}, f"第{index + 1}封新人草稿To错误"),
            (not _addresses(message, "Cc") and not _addresses(message, "Bcc"), f"第{index + 1}封新人草稿含Cc/Bcc"),
            (str(message.get("Subject", "")) == subject, f"第{index + 1}封新人草稿主题错误"),
            (str(message.get("Date", "")) == "Wed, 15 Apr 2026 09:00:00 +0000", f"第{index + 1}封新人草稿日期错误"),
            (str(message.get("Message-ID", "")) == message_id, f"第{index + 1}封新人草稿Message-ID错误"),
            (str(message.get("X-EvoCowork-State", "")) == "DRAFT_NOT_SENT" and str(message.get("X-EvoCowork-Type", "")) == "NEW_HIRE_BUDDY_NOTIFICATION", f"第{index + 1}封新人草稿状态或类型错误"),
            (_body(message) == body, f"第{index + 1}封新人草稿姓名、Buddy、部门或开始日错误"),
        ])
    return _ratio(findings)


def _dm_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    observed = json.loads((artifact / DM_JSON).read_text(encoding="utf-8"))
    return _ratio(_leaf_findings(observed, DM_EXPECTED, "Buddy_Slack_DM_Drafts"))


def _head_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    messages = _messages(artifact / HEAD_MBOX)
    findings: list[tuple[bool, str]] = [(len(messages) == 1, "部门负责人升级MBOX不是恰好一封")]
    if messages:
        message = messages[0]
        expected_body = (
            "Nora Lindgren starts 2026-04-21 in Sales. Brandon Osei and Ava Chen do not meet "
            "the one-year tenure rule, so no eligible same-department buddy is available. "
            "Please identify an eligible Sales buddy manually. No cross-department buddy was assigned. DRAFT_NOT_SENT."
        )
        findings.extend([
            (_addresses(message, "From") == {"hr-ops@company.example.com"}, "升级草稿From错误"),
            (_addresses(message, "To") == {"vp-sales@company.example.com"}, "升级草稿未只发VP Sales"),
            (not _addresses(message, "Cc") and not _addresses(message, "Bcc"), "升级草稿含Cc/Bcc"),
            (str(message.get("Subject", "")) == "Manual buddy needed for Nora Lindgren", "升级草稿主题错误"),
            (str(message.get("Date", "")) == "Wed, 15 Apr 2026 09:00:00 +0000", "升级草稿日期错误"),
            (str(message.get("Message-ID", "")) == "<buddy-head-5116-01@draft.invalid>", "升级草稿Message-ID错误"),
            (str(message.get("X-EvoCowork-State", "")) == "DRAFT_NOT_SENT" and str(message.get("X-EvoCowork-Type", "")) == "DEPARTMENT_HEAD_BUDDY_ESCALATION", "升级草稿状态或类型错误"),
            (_body(message) == expected_body, "升级草稿Nora、Brandon、Ava、年资或跨部门说明错误"),
        ])
    else:
        findings.extend((False, f"缺少升级草稿字段{index}") for index in range(8))
    return _ratio(findings)


def _execution_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    observed = json.loads((artifact / EXECUTION_JSON).read_text(encoding="utf-8"))
    return _ratio(_leaf_findings(observed, EXECUTION_EXPECTED, "Buddy_Execution_Control"))


def _word_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    document = Document(artifact / DOCX)
    rows = [tuple(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
    text = "\n".join(p.text for p in document.paragraphs) + "\n" + "\n".join(" | ".join(row) for row in rows)
    findings: list[tuple[bool, str]] = [(len(document.sections) == 3 and len(document.tables) == 3, "Word不是三页三表")]
    for expected in [
        ("Li Wei", "Engineering", "ASSIGN_BUDDY", "Carlos Mendoza"),
        ("Li Wei-Chen", "Marketing", "ASSIGN_BUDDY", "Sandra Liu"),
        ("Nora Lindgren", "Sales", "ESCALATE_DEPARTMENT_HEAD", "vp-sales@company.example.com"),
        ("Carlos Mendoza", "Engineering", "2023-03-20", "ELIGIBLE"),
        ("Fatima Al-Rashid", "Engineering", "2024-06-15", "ON_LEAVE"),
        ("James Whitfield", "Engineering", "2025-08-01", "TENURE_UNDER_ONE_YEAR"),
        ("Sandra Liu", "Marketing", "2022-09-10", "ELIGIBLE"),
        ("Olga Petrova", "Marketing", "2025-06-20", "TENURE_UNDER_ONE_YEAR"),
        ("Brandon Osei", "Sales", "2025-11-01", "TENURE_UNDER_ONE_YEAR"),
        ("Ava Chen", "Sales", "2026-01-15", "TENURE_UNDER_ONE_YEAR"),
        ("New hire emails", "2", "0", "DRAFT_NOT_SENT"),
        ("Buddy Slack DMs", "2", "0", "DRAFT_NOT_DELIVERED"),
        ("Department head escalation", "1", "0", "DRAFT_NOT_SENT"),
        ("Assignments", "2", "0", "DRAFT_NOT_APPLIED"),
        ("Real side effects", "0", "0", "NONE"),
    ]:
        findings.append((expected in rows, f"Word缺少任务专属行：{expected[0]}"))
    for token in ("Three next-week hires", "Fatima is unavailable", "under one year of tenure", "No cross-department substitution", "offline"):
        findings.append((token.casefold() in text.casefold(), f"Word缺少 {token}"))
    findings.append(("nora was assigned a cross-department buddy" not in text.casefold(), "Word虚假声称Nora跨部门配对"))
    return _ratio(findings)


def _ppt_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    deck = Presentation(artifact / PPTX)
    slides: list[str] = []
    rows_by_slide: list[list[tuple[str, ...]]] = []
    table_counts: list[int] = []
    pictures = 0
    for slide in deck.slides:
        text_parts: list[str] = []
        rows: list[tuple[str, ...]] = []
        tables = 0
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                text_parts.append(shape.text)
            if getattr(shape, "has_table", False):
                tables += 1
                rows.extend(tuple(cell.text for cell in row.cells) for row in shape.table.rows)
        slides.append("\n".join(text_parts))
        rows_by_slide.append(rows)
        table_counts.append(tables)
    findings: list[tuple[bool, str]] = [
        (len(deck.slides) == 3, "PPT不是三页"),
        (table_counts == [1, 1, 1], "PPT不是每页一张原生表"),
        (pictures == 0, "PPT使用图片替代关键内容"),
    ]
    expected_rows = {
        0: [
            ("Li Wei", "Engineering", "ASSIGN_BUDDY", "Carlos Mendoza"),
            ("Li Wei-Chen", "Marketing", "ASSIGN_BUDDY", "Sandra Liu"),
            ("Nora Lindgren", "Sales", "ESCALATE_DEPARTMENT_HEAD", "VP Sales"),
        ],
        1: [
            ("Carlos Mendoza", "Engineering", "2023-03-20", "ELIGIBLE"),
            ("Fatima Al-Rashid", "Engineering", "2024-06-15", "ON_LEAVE"),
            ("James Whitfield", "Engineering", "2025-08-01", "TENURE_UNDER_ONE_YEAR"),
            ("Sandra Liu", "Marketing", "2022-09-10", "ELIGIBLE"),
            ("Olga Petrova", "Marketing", "2025-06-20", "TENURE_UNDER_ONE_YEAR"),
            ("Brandon Osei", "Sales", "2025-11-01", "TENURE_UNDER_ONE_YEAR"),
            ("Ava Chen", "Sales", "2026-01-15", "TENURE_UNDER_ONE_YEAR"),
        ],
        2: [
            ("New hire email", "2", "0", "DRAFT_NOT_SENT"),
            ("Buddy Slack DM", "2", "0", "DRAFT_NOT_DELIVERED"),
            ("Head escalation", "1", "0", "DRAFT_NOT_SENT"),
            ("Assignments", "2", "0", "NOT_APPLIED"),
        ],
    }
    for slide_index, expected in expected_rows.items():
        rows = rows_by_slide[slide_index] if slide_index < len(rows_by_slide) else []
        for row in expected:
            findings.append((row in rows, f"PPT第{slide_index + 1}页缺少 {row[0]} 行"))
    for slide_index, tokens in enumerate([
        ("Buddy assignment outcome", "3 new hires", "2 assigned", "1 department-head escalation", "0 cross-department"),
        ("Policy filters the pool", "Same department", "before April 2025", "available", "longest tenure"),
        ("All actions remain drafts", "Nora receives no buddy email", "Sales supplies an eligible buddy"),
    ]):
        text = slides[slide_index] if slide_index < len(slides) else ""
        for token in tokens:
            findings.append((token in text, f"PPT第{slide_index + 1}页缺少 {token}"))
    return _ratio(findings)


def _office_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    book = load_workbook(artifact / XLSX, data_only=False)
    table_count = sum(len(sheet.tables) for sheet in book.worksheets)
    chart_count = sum(len(sheet._charts) for sheet in book.worksheets)
    conditional_count = sum(len(sheet.conditional_formatting) for sheet in book.worksheets)
    sheet_count = len(book.sheetnames)
    book.close()
    document = Document(artifact / DOCX)
    deck = Presentation(artifact / PPTX)
    ppt_tables = [sum(getattr(shape, "has_table", False) for shape in slide.shapes) for slide in deck.slides]
    ppt_pictures = sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for slide in deck.slides for shape in slide.shapes)
    findings = [
        (sheet_count == 12 and table_count == 12, "Excel不是十二表十二原生表"),
        (chart_count >= 1, "Excel缺少原生图表"),
        (conditional_count >= 1, "Excel缺少候选资格条件格式"),
        (len(document.sections) == 3 and len(document.tables) == 3, "Word三页三表结构错误"),
        (len(deck.slides) == 3 and ppt_tables == [1, 1, 1] and ppt_pictures == 0, "PPT三页三表或图片替代结构错误"),
    ]
    for name in (XLSX, DOCX, PPTX):
        safe, detail = V1.package_safe(artifact / name)
        findings.append((safe, f"{name} 含宏、外链或危险对象：{detail}"))
    return _ratio(findings)


def _assignment_severity(artifact: Path) -> tuple[float, list[str], float]:
    book = load_workbook(artifact / XLSX, data_only=True)
    pairs = _rows(book["Assigned_Pairs"])
    unfilled = _rows(book["Unfilled_Assignments"])
    book.close()
    severities: list[tuple[float, str]] = []
    banned = {"Fatima Al-Rashid", "James Whitfield", "Olga Petrova", "Brandon Osei", "Ava Chen"}
    if any(len(row) > 2 and row[2] in banned for row in pairs):
        severities.append((0.85, "分配了休假或年资不足候选"))
    departments = {"Li Wei": "Engineering", "Li Wei-Chen": "Marketing", "Nora Lindgren": "Sales"}
    buddy_departments = {row[1]: row[3] for row in POOL}
    if any(len(row) > 2 and departments.get(row[0]) != buddy_departments.get(row[2]) for row in pairs):
        severities.append((0.95, "发生跨部门Buddy分配"))
    buddies = [row[2] for row in pairs if len(row) > 2]
    if len(buddies) != len(set(buddies)):
        severities.append((0.75, "同一Buddy被重复分配"))
    if any(row and row[0] == "Nora Lindgren" for row in pairs) or not any(row and row[0] == "Nora Lindgren" for row in unfilled):
        severities.append((0.9, "Nora缺口被错误处理或未升级"))
    if pairs != PAIRS:
        severities.append((0.35, "两对指定分配不准确"))
    score = max((value for value, _ in severities), default=0.0)
    return score, [label for _, label in severities], 1.0


def _execution_severity(artifact: Path) -> tuple[float, list[str], float]:
    new_hire = _messages(artifact / NEW_HIRE_MBOX)
    head = _messages(artifact / HEAD_MBOX)
    dm = json.loads((artifact / DM_JSON).read_text(encoding="utf-8"))
    execution = json.loads((artifact / EXECUTION_JSON).read_text(encoding="utf-8"))
    severities: list[tuple[float, str]] = []
    allowed_new_hires = ["li.wei@company.example.com", "li.weichen@company.example.com"]
    if any(_addresses(message, "To") != {wanted} for message, wanted in zip(new_hire, allowed_new_hires)):
        severities.append((0.45, "新人邮件收件人越权或错误"))
    if head and _addresses(head[0], "To") != {"vp-sales@company.example.com"}:
        severities.append((0.5, "部门负责人升级收件人错误"))
    if any(str(message.get("X-EvoCowork-State", "")) != "DRAFT_NOT_SENT" for message in (*new_hire, *head)):
        severities.append((0.85, "邮件被标记为已发送"))
    if dm.get("delivered_count") not in (0, None) or dm.get("state") not in ("DRAFT_NOT_DELIVERED", None):
        severities.append((0.85, "Buddy DM被标记为已交付"))
    if execution.get("emails_sent") not in (0, None) or execution.get("buddy_dms_delivered") not in (0, None):
        severities.append((0.85, "Execution声称邮件或DM已执行"))
    if execution.get("assignments_applied") not in (0, None):
        severities.append((0.9, "Execution声称Buddy分配已应用"))
    if execution.get("cross_department_assignments") not in (0, None):
        severities.append((0.95, "Execution声称跨部门分配"))
    if execution.get("real_side_effects") not in (0, None):
        severities.append((1.0, "Execution声称真实副作用"))
    score = max((value for value, _ in severities), default=0.0)
    return score, [label for _, label in severities], 1.0


def _extra_severity(artifact: Path) -> tuple[float, list[str], float]:
    files = {path.name for path in artifact.iterdir() if path.is_file()}
    directories = [path.name for path in artifact.iterdir() if path.is_dir()]
    extras = sorted(files - set(REQUIRED))
    active = bool(extras or directories)
    return 1.0 if active else 0.0, [f"输出含额外文件或子目录：{extras + directories}"] if active else [], 1.0


def _safe_positive(call: Callable[[Path], tuple[float, list[str], int, int]], artifact: Path) -> tuple[float, list[str], float, float]:
    try:
        return call(artifact)
    except Exception as exc:
        return 0.0, [f"任务专属检查无法完成：{type(exc).__name__}: {exc}"], 0, 1


def _safe_penalty(call: Callable[[Path], tuple[float, list[str], float]], artifact: Path) -> tuple[float, list[str], float]:
    try:
        return call(artifact)
    except Exception as exc:
        return 1.0, [f"严重边界无法核验：{type(exc).__name__}: {exc}"], 1.0


def evaluate(artifact: Path) -> dict[str, Any]:
    ok, errors = _core(artifact)
    hashes = [_sha256(artifact / name) for name in REQUIRED if (artifact / name).is_file()]
    if not ok:
        return _fatal(errors, hashes)
    legacy = _legacy_exact(artifact)
    legacy_checks = legacy.get("criteria", {})
    custom = {
        "R002": _source_positive, "R003": _formula_positive, "R004": _eligibility_positive,
        "R005": _assignment_positive, "R006": _new_hire_email_positive, "R007": _dm_positive,
        "R008": _head_positive, "R009": _execution_positive, "R010": _word_positive,
        "R011": _ppt_positive, "R012": _office_positive,
    }
    legacy_map = {
        "R002": ("R002",), "R003": ("R003", "R004"), "R004": ("R005",),
        "R005": ("R005", "R006"), "R006": ("R007",), "R007": ("R008",),
        "R008": ("R008",), "R009": ("R009",), "R010": ("R010",),
        "R011": ("R011",), "R012": ("R012",),
    }
    results = [_criterion("R001", 1, "七份核心交付存在、非空且可解析", 7, 7)]
    scores: dict[str, float] = {"R001": 1.0}
    for key in [f"R{index:03d}" for index in range(2, 13)]:
        score, failures, earned, total = _safe_positive(custom[key], artifact)
        for legacy_key in legacy_map[key]:
            check = legacy_checks.get(legacy_key, {"passed": False, "details": ["v1 未返回检查"]})
            if score == 1 and not check.get("passed"):
                total += 1
                failures.extend(str(item) for item in (check.get("details") or [f"v1 {legacy_key} 检查失败"]))
                score = earned / total
        results.append(_criterion(key, score, "本题任务专属检查通过" if not failures else "未通过：" + "；".join(failures), earned, total))
        scores[key] = score
    for key, call in (("P001", _assignment_severity), ("P002", _execution_severity), ("P003", _extra_severity)):
        score, labels, total = _safe_penalty(call, artifact)
        results.append(_criterion(key, score, "未触发任务专属严重边界" if not labels else "；".join(labels), score * total, total))
        scores[key] = score
    positive_total = sum(POSITIVE_WEIGHTS.values())
    positive_earned = sum(POSITIVE_WEIGHTS[key] * scores[key] for key in POSITIVE_WEIGHTS)
    penalty_points = sum(abs(PENALTY_WEIGHTS[key]) * scores[key] for key in PENALTY_WEIGHTS)
    reward = round(max(0.0, min(1.0, (positive_earned - penalty_points) / positive_total)), 6)
    return {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes, "criterion_results": results,
        "reward": reward, "passed": reward >= 0.8, "hurdle_failures": [],
        "penalties": [{"id": key, "amount": round(abs(PENALTY_WEIGHTS[key]) * scores[key] / positive_total, 6)} for key in PENALTY_WEIGHTS if scores[key] > 0],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    parser.add_argument("--scoring-version", choices=("v1", "v2"), default="v2")
    args = parser.parse_args()
    try:
        result = V1.evaluate(args.artifact_dir) if args.scoring_version == "v1" else evaluate(args.artifact_dir)
    except Exception as exc:
        result = _fatal([f"Judge异常：{type(exc).__name__}: {exc}"], [])
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{float(result['reward']):.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
