#!/usr/bin/env python3
"""AutomationBench 5116 的任务专属 v2 分层负例。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from typing import Callable

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
INPUTS = ROOT / "environment" / "input_files"
OUTPUTS = (
    "Buddy_Assignment_Audit.xlsx",
    "Buddy_Assignment_Brief.docx",
    "Buddy_Assignment_Review.pptx",
    "New_Hire_Buddy_Drafts.mbox",
    "Buddy_Slack_DM_Drafts.json",
    "Department_Head_Escalation_Draft.mbox",
    "Buddy_Execution_Control.json",
)
Mutator = Callable[[Path], None]


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"无法载入 {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _xlsx(root: Path, sheet: str, cell: str, value) -> None:
    path = root / OUTPUTS[0]
    book = load_workbook(path)
    book[sheet][cell] = value
    book.save(path)


def _json(root: Path, name: str, mutator: Callable[[dict], None]) -> None:
    path = root / name
    payload = json.loads(path.read_text(encoding="utf-8"))
    mutator(payload)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _rewrite_mbox(root: Path, filename: str, mutator: Callable[[list], None]) -> None:
    path = root / filename
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    mutator(messages)
    with path.open("wb") as handle:
        for index, message in enumerate(messages):
            handle.write(f"From MAILER-DAEMON Wed Apr 15 09:00:{index:02d} 2026\n".encode())
            BytesGenerator(handle, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
            handle.write(b"\n")


def _word_false(root: Path) -> None:
    path = root / OUTPUTS[1]
    document = Document(path)
    document.add_paragraph("Nora was assigned a cross-department buddy.")
    document.save(path)


def _ppt_omit_nora(root: Path) -> None:
    path = root / OUTPUTS[2]
    deck = Presentation(path)
    for slide in deck.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                shape.text = shape.text.replace("Nora Lindgren", "Unknown")
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        cell.text = cell.text.replace("Nora Lindgren", "Unknown")
    deck.save(path)


def _input_only(root: Path) -> None:
    for path in list(root.iterdir()):
        if path.is_file():
            path.unlink()
    for source in INPUTS.iterdir():
        if source.is_file():
            shutil.copy2(source, root / source.name)


def _cases() -> list[tuple[str, Mutator, str]]:
    cases: list[tuple[str, Mutator, str]] = []
    for name in OUTPUTS:
        cases.append((f"missing-{name}", lambda root, n=name: (root / n).unlink(), "fatal_zero"))
        cases.append((f"corrupt-{name}", lambda root, n=name: (root / n).write_bytes(b"not-a-valid-artifact"), "fatal_zero"))
    cases.extend([
        ("new-hire-mbox-empty", lambda root: (root / OUTPUTS[3]).write_bytes(b""), "fatal_zero"),
        ("unmodified-input-only", _input_only, "fatal_zero"),
        ("fatima-assigned", lambda root: _xlsx(root, "Assigned_Pairs", "C2", "Fatima Al-Rashid"), "partial"),
        ("li-wei-buddy-swapped", lambda root: _xlsx(root, "Assigned_Pairs", "C2", "Sandra Liu"), "partial"),
        ("nora-cross-department", lambda root: _xlsx(root, "Unfilled_Assignments", "A2", "Li Wei"), "partial"),
        ("assignment-formula-removed", lambda root: _xlsx(root, "Assignment_Decisions", "J2", "ASSIGN_BUDDY"), "partial"),
        ("wrong-new-hire-recipient", lambda root: _rewrite_mbox(root, OUTPUTS[3], lambda ms: ms[0].replace_header("To", "nora.lindgren@company.example.com")), "partial"),
        ("new-hire-email-sent", lambda root: _rewrite_mbox(root, OUTPUTS[3], lambda ms: ms[0].replace_header("X-EvoCowork-State", "SENT")), "partial"),
        ("wrong-dm-recipient", lambda root: _json(root, OUTPUTS[4], lambda p: p["drafts"][0].update({"recipient_id": "U_SANDRA"})), "partial"),
        ("dm-delivered", lambda root: _json(root, OUTPUTS[4], lambda p: p.update({"delivered_count": 2, "state": "DELIVERED"})), "partial"),
        ("wrong-head-recipient", lambda root: _rewrite_mbox(root, OUTPUTS[5], lambda ms: ms[0].replace_header("To", "vp-eng@company.example.com")), "partial"),
        ("nora-false-buddy-notification", lambda root: _rewrite_mbox(root, OUTPUTS[3], lambda ms: ms[0].set_payload(ms[0].get_payload() + "\nNora Lindgren has an assigned buddy.")), "partial"),
        ("assignments-applied", lambda root: _json(root, OUTPUTS[6], lambda p: p.update({"assignments_applied": 2})), "partial"),
        ("cross-department-count", lambda root: _json(root, OUTPUTS[6], lambda p: p.update({"cross_department_assignments": 1})), "partial"),
        ("word-false-cross-department", _word_false, "partial"),
        ("ppt-omit-nora", _ppt_omit_nora, "partial"),
        ("extra-file", lambda root: (root / "unexpected.txt").write_text("unexpected", encoding="utf-8"), "partial"),
    ])
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    judge = _load("buddy_5116_v2_judge", ROOT / "tests" / "judge.py")
    judge.V1.render = lambda path: ({".xlsx": 12, ".docx": 3, ".pptx": 3}[path.suffix.casefold()], "render-ok")
    if judge.evaluate(GOLDEN).get("reward") != 1.0:
        raise AssertionError("Golden v2 必须为1.0")
    results: list[dict] = []
    for name, mutator, layer in _cases():
        with tempfile.TemporaryDirectory(prefix=f"buddy-5116-v2-{name}-") as temporary:
            output = Path(temporary) / "output"
            shutil.copytree(GOLDEN, output)
            mutator(output)
            result = judge.evaluate(output)
            reward = float(result.get("reward", 0))
            if layer == "fatal_zero" and reward != 0:
                raise AssertionError(f"{name} 应硬归零，实际 {reward}")
            if layer == "partial" and not (0 < reward < 1):
                raise AssertionError(f"{name} 应保留连续部分分，实际 {reward}")
            results.append({"case": name, "layer": layer, "reward": reward, "passed": bool(result.get("passed")), "penalties": result.get("penalties", [])})
    partial = [item["reward"] for item in results if item["layer"] == "partial"]
    report = {
        "total": len(results),
        "fatal_zero": sum(item["layer"] == "fatal_zero" for item in results),
        "partial_continuous": len(partial),
        "partial_reward_range": [min(partial), max(partial)],
        "cases": results,
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
