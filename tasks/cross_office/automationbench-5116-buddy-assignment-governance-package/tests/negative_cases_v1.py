#!/usr/bin/env python3
import json,mailbox,shutil,subprocess,sys,tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1];GOLDEN=ROOT/"solution/Golden";PYTHON=sys.executable
def score(case):
 result=case.parent/(case.name+".json");reward=case.parent/(case.name+".txt");subprocess.run([PYTHON,str(ROOT/"tests/judge.py"),"--artifact-dir",str(case),"--result",str(result),"--reward-file",str(reward)],check=True,capture_output=True,text=True,timeout=240);return float(reward.read_text())
def rewrite(path,mutation):
 box=mailbox.mbox(path,create=False);messages=list(box);box.close();mutation(messages)
 with path.open("wb") as handle:
  for index,message in enumerate(messages):handle.write(f"From x Wed Apr 15 09:00:{index:02d} 2026\n".encode());BytesGenerator(handle,policy=SMTP,mangle_from_=True,maxheaderlen=0).flatten(message);handle.write(b"\n")
def main():
 with tempfile.TemporaryDirectory(prefix="buddy5116-neg-") as temp:
  root=Path(temp);names=["missing","wrong_pairs","fatima_assigned","cross_department","wrong_newhire","wrong_dm","wrong_head","nora_notified","emails_sent","applied","formula_removed","word_false","ppt_omit"];cases=[]
  for name in names:case=root/name;shutil.copytree(GOLDEN,case);cases.append(case)
  (cases[0]/"Buddy_Assignment_Review.pptx").unlink()
  payload=json.loads((cases[1]/"Buddy_Execution_Control.json").read_text());payload["assigned_pairs"]=3;(cases[1]/"Buddy_Execution_Control.json").write_text(json.dumps(payload))
  rewrite(cases[2]/"New_Hire_Buddy_Drafts.mbox",lambda ms:ms[0].set_payload(ms[0].get_payload()+"\nFatima Al-Rashid"))
  payload=json.loads((cases[3]/"Buddy_Execution_Control.json").read_text());payload["cross_department_assignments"]=1;(cases[3]/"Buddy_Execution_Control.json").write_text(json.dumps(payload))
  rewrite(cases[4]/"New_Hire_Buddy_Drafts.mbox",lambda ms:ms[0].replace_header("To","nora.lindgren@company.example.com"))
  payload=json.loads((cases[5]/"Buddy_Slack_DM_Drafts.json").read_text());payload["drafts"][0]["recipient_id"]="U_SANDRA";(cases[5]/"Buddy_Slack_DM_Drafts.json").write_text(json.dumps(payload))
  rewrite(cases[6]/"Department_Head_Escalation_Draft.mbox",lambda ms:ms[0].replace_header("To","vp-eng@company.example.com"))
  rewrite(cases[7]/"New_Hire_Buddy_Drafts.mbox",lambda ms:ms[0].set_payload(ms[0].get_payload()+"\nNora Lindgren"))
  payload=json.loads((cases[8]/"Buddy_Execution_Control.json").read_text());payload["emails_sent"]=3;(cases[8]/"Buddy_Execution_Control.json").write_text(json.dumps(payload))
  payload=json.loads((cases[9]/"Buddy_Execution_Control.json").read_text());payload["assignments_applied"]=2;(cases[9]/"Buddy_Execution_Control.json").write_text(json.dumps(payload))
  workbook=load_workbook(cases[10]/"Buddy_Assignment_Audit.xlsx");workbook["Assignment_Decisions"]["J2"]="ASSIGN_BUDDY";workbook.save(cases[10]/"Buddy_Assignment_Audit.xlsx")
  document=Document(cases[11]/"Buddy_Assignment_Brief.docx");document.add_paragraph("Nora was assigned a cross-department buddy.");document.save(cases[11]/"Buddy_Assignment_Brief.docx")
  presentation=Presentation(cases[12]/"Buddy_Assignment_Review.pptx")
  for slide in presentation.slides:
   for shape in slide.shapes:
    if getattr(shape,"has_text_frame",False):shape.text=shape.text.replace("Nora Lindgren","Unknown")
    if getattr(shape,"has_table",False):
     for row in shape.table.rows:
      for cell in row.cells:cell.text=cell.text.replace("Nora Lindgren","Unknown")
  presentation.save(cases[12]/"Buddy_Assignment_Review.pptx")
  rewards={case.name:score(case) for case in cases};print(json.dumps(rewards,indent=2));assert all(value==0 for value in rewards.values()),rewards
if __name__=="__main__":main()
