#!/usr/bin/env python3
"""只把 AutomationBench 5117 的公开初始状态静态化为本题输入。

本脚本只接受 ``hr.grievance_routing``，逐项断言八个申诉案件、两封权威
程序邮件、冻结时间、Slack 频道和九条上游断言。它不是通用制题器，不会
创造案件、人员、严重度、申诉内容、路由规则或收件地址。
"""

from __future__ import annotations

import hashlib
import importlib.util
import json
import shutil
import sys
import types
import zipfile
from datetime import datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
UPSTREAM = Path(__file__).with_name("upstream-hr-tasks.py")
SNAPSHOT = Path(__file__).with_name("upstream-task-5117.json")
LICENSE = Path(__file__).with_name("LICENSE-AutomationBench-MIT")
DATA_DIR = ROOT / "data" / "input_files"
ENV_DIR = ROOT / "environment" / "input_files"
FIXED_ZIP_TIME = (2026, 4, 15, 9, 0, 0)
SOURCE_SHA256 = "c173a36a03f43096bd122577584aa32f9df379d8760aa46c43be04c85d101779"
LICENSE_SHA256 = "c6f9d87aff6d627d5cd78cb09364dafed7d6f1c4bb7a91183d57377157e96ba4"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_task() -> dict:
    if sha256(UPSTREAM) != SOURCE_SHA256 or sha256(LICENSE) != LICENSE_SHA256:
        raise ValueError("AutomationBench 冻结源文件或许可证哈希错误")
    datasets = types.ModuleType("datasets")
    datasets.Dataset = object
    sys.modules["datasets"] = datasets
    for name in ("automationbench", "automationbench.domains", "automationbench.domains.hr"):
        sys.modules[name] = types.ModuleType(name)
    noise = types.ModuleType("automationbench.domains.hr._noise")
    noise.apply_noise = lambda value: value
    sys.modules[noise.__name__] = noise
    spec = importlib.util.spec_from_file_location("automationbench_hr_tasks_frozen_5117", UPSTREAM)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入冻结的 AutomationBench HR 源文件")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    task = module.get_task_5117()
    if task.get("example_id") != 5117 or task.get("task") != "hr.grievance_routing":
        raise ValueError("冻结函数不再是 AutomationBench 5117")
    return task


def assert_shape(task: dict) -> None:
    prompt = task["prompt"][1]["content"]
    assert all(token in prompt for token in ("pending employee grievances", "Update each case status", "notification emails", "Google Sheets and Gmail"))
    state = task["info"]["initial_state"]
    assert state["meta"] == {"schema_version": "0.1.0", "current_time": "2026-04-15T09:00:00Z"}
    messages = state["gmail"]["messages"]
    assert [message["id"] for message in messages] == ["msg_grievance_handbook_5117", "msg_griev_override_5117"]
    handbook = messages[0]["body_plain"]
    assert all(token in handbook for token in ("Level 1", "Level 2", "Level 3", "strictly confidential", "Do not post", "manager IS the subject", "Routed to [handler]"))
    override = messages[1]["body_plain"]
    assert all(token in override for token in ("ALL discrimination-category grievances", "Legal AND HR Director", "overrides"))
    spreadsheet = state["google_sheets"]["spreadsheets"][0]
    assert spreadsheet["id"] == "ss_grievances_5117" and spreadsheet["title"] == "Grievance Tracker"
    assert [worksheet["title"] for worksheet in spreadsheet["worksheets"]] == ["Pending Cases"]
    rows = spreadsheet["worksheets"][0]["rows"]
    assert [row["row_id"] for row in rows] == list(range(2, 10))
    assert [row["cells"]["Case ID"] for row in rows] == [f"GR-2026-0{number}" for number in range(41, 49)]
    assert [row["cells"]["Employee"] for row in rows] == ["Anna Svensson", "Ben Torres", "Carmen Obi", "David Holm", "Elise Braun", "Frank Yoshida", "Grace Petrov", "Hugo Strand"]
    assert rows[2]["cells"]["Category"] == "Discrimination" and rows[2]["cells"]["Severity"] == "Level 1"
    assert rows[4]["cells"]["Notes"] == "Complaint is about direct manager"
    assert rows[6]["cells"]["Category"] == "Discrimination" and rows[6]["cells"]["Severity"] == "Level 2"
    assert state["slack"] == {"channels": [{"id": "C_HR_5117", "name": "hr-general", "is_private": False}], "messages": []}
    assertions = task["info"]["assertions"]
    assert len(assertions) == 9
    assertion_text = json.dumps(assertions)
    assert all(token in assertion_text for token in ("GR-2026-041", "GR-2026-042", "GR-2026-043", "GR-2026-045", "slack_message_not_exists"))
    # 上游断言只抽查四案，本题 Judge 必须按全表和政策覆盖八案。
    assert all(case not in assertion_text for case in ("GR-2026-044", "GR-2026-046", "GR-2026-047", "GR-2026-048"))


def normalize_xlsx(path: Path) -> None:
    temporary = path.with_suffix(".normalized.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    temporary.replace(path)


def style_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    widths = {"A": 12, "B": 21, "C": 21, "D": 34, "E": 22, "F": 15, "G": 48, "H": 18, "I": 34}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=8)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = sheet.dimensions


def write_tracker(state: dict, target: Path) -> None:
    source = state["google_sheets"]["spreadsheets"][0]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Pending Cases"
    workbook.creator = "AutomationBench 5117 public fixture"
    workbook.title = source["title"]
    workbook.subject = "Static Office representation of eight pending grievance cases"
    workbook.description = "No case, employee, manager, category, severity, description, status or note was added beyond AutomationBench 5117."
    headers = ["Source Row", "Case ID", "Employee", "Manager", "Manager Email", "Category", "Severity", "Description", "Status", "Notes"]
    sheet.append(headers)
    for item in source["worksheets"][0]["rows"]:
        cells = item["cells"]
        sheet.append([item["row_id"], cells["Case ID"], cells["Employee"], cells["Manager"], cells["Manager Email"], cells["Category"], cells["Severity"], cells["Description"], cells["Status"], cells["Notes"]])
    table = Table(displayName="PendingGrievanceSourceInput", ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    style_sheet(sheet)
    workbook.save(target)
    normalize_xlsx(target)


def message_bytes(record: dict) -> bytes:
    message = EmailMessage(policy=SMTP)
    message["From"] = record["from_"]
    message["To"] = ", ".join(record["to"])
    message["Subject"] = record["subject"]
    parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
    message["Date"] = parsed.strftime("%a, %d %b %Y %H:%M:%S +0000")
    message["Message-ID"] = f"<{record['id']}@automationbench.invalid>"
    message["X-AutomationBench-Message-ID"] = record["id"]
    message["X-AutomationBench-Thread-ID"] = record["thread_id"]
    message["X-AutomationBench-Labels"] = ",".join(record.get("label_ids", []))
    message.set_content(record["body_plain"], charset="utf-8")
    buffer = BytesIO()
    BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
    return buffer.getvalue()


def write_mbox(messages: list[dict], target: Path) -> None:
    chunks: list[bytes] = []
    for record in messages:
        parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
        chunks.append(f"From {record['from_']} {parsed.strftime('%a %b %d %H:%M:%S %Y')}\n".encode("ascii"))
        chunks.append(message_bytes(record).replace(b"\r\n", b"\n"))
        chunks.append(b"\n")
    target.write_bytes(b"".join(chunks))


def main() -> None:
    task = load_task()
    assert_shape(task)
    state = task["info"]["initial_state"]
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    ENV_DIR.mkdir(parents=True, exist_ok=True)
    SNAPSHOT.write_text(json.dumps(task, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_tracker(state, DATA_DIR / "Grievance_Tracker.xlsx")
    write_mbox(state["gmail"]["messages"], DATA_DIR / "Grievance_Procedures.mbox")
    (DATA_DIR / "Slack_Channels.json").write_text(json.dumps(state["slack"], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (DATA_DIR / "Workflow_Context.json").write_text(json.dumps({"meta": state["meta"]}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    for source in sorted(DATA_DIR.iterdir()):
        if source.is_file():
            shutil.copy2(source, ENV_DIR / source.name)


if __name__ == "__main__":
    main()
