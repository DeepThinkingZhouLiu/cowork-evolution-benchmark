#!/usr/bin/env python3
"""AutomationBench 5117 申诉路由与保密控制包的逐题参考解法。"""

from __future__ import annotations

import argparse
import hashlib
import json
import mailbox
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt as PptPt


INPUT_HASHES = {
    "Grievance_Tracker.xlsx": "65cb8f8c648072bc7a11cb265dbe63e5d1c1c2c691049593ffa7f6f252ba8ae2",
    "Grievance_Procedures.mbox": "a95ec8ee29877e69756b51715672a8f1ed736a00766eef326879e53eab518691",
    "Slack_Channels.json": "f5be48025181a0cae840271920c0998a2b28068c874bbf8498f3b44720da45ba",
    "Workflow_Context.json": "76c4a0657bc27f2fd830b10613ec7ceae3adcdc5798aeee0f822d27bf69bbc99",
}
XLSX_NAME = "Grievance_Routing_Audit.xlsx"
DOCX_NAME = "Grievance_Routing_Control_Brief.docx"
PPTX_NAME = "Grievance_Routing_Review.pptx"
MBOX_NAME = "Grievance_Notification_Drafts.mbox"
UPDATES_NAME = "Grievance_Tracker_Update_Drafts.json"
SLACK_NAME = "Grievance_Slack_Control.json"
EXECUTION_NAME = "Grievance_Execution_Control.json"
FIXED_ISO_TIME = "2026-04-15T09:00:00Z"
FIXED_ZIP_TIME = (2026, 4, 15, 9, 0, 0)

NAVY = "17365D"; BLUE = "5B9BD5"; GREEN = "E2F0D9"; AMBER = "FFF2CC"; RED = "F4CCCC"; GRAY = "E7E6E6"; WHITE = "FFFFFF"
THIN = Side(style="thin", color="B7C9D6")
PPT_NAVY = PptRGBColor(23, 54, 93); PPT_BLUE = PptRGBColor(91, 155, 213); PPT_LIGHT_BLUE = PptRGBColor(221, 235, 247)
PPT_GREEN = PptRGBColor(226, 240, 217); PPT_AMBER = PptRGBColor(255, 242, 204); PPT_RED = PptRGBColor(244, 204, 204); PPT_WHITE = PptRGBColor(255, 255, 255)


@dataclass(frozen=True)
class GrievanceCase:
    source_row: int
    case_id: str
    employee: str
    manager: str
    manager_email: str
    category: str
    severity: str
    description: str
    status: str
    notes: str


@dataclass(frozen=True)
class Decision:
    source: GrievanceCase
    discrimination_override: str
    manager_subject: str
    route_type: str
    handler: str
    to_email: str
    cc_email: str
    proposed_status: str
    authority: str


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def decode_message(message) -> str:
    if message.is_multipart():
        return "\n".join((part.get_payload(decode=True) or b"").decode(part.get_content_charset() or "utf-8", errors="replace") for part in message.walk() if part.get_content_type() == "text/plain").strip()
    return (message.get_payload(decode=True) or b"").decode(message.get_content_charset() or "utf-8", errors="replace").strip()


def read_inputs(input_dir: Path) -> tuple[list[GrievanceCase], list[dict], dict, str]:
    for name, expected in INPUT_HASHES.items():
        source = input_dir / name
        if not source.is_file() or sha256(source) != expected:
            raise ValueError(f"输入 {name} 缺失或 SHA256 与冻结版本不一致")
    context = json.loads((input_dir / "Workflow_Context.json").read_text(encoding="utf-8"))
    if context != {"meta": {"schema_version": "0.1.0", "current_time": FIXED_ISO_TIME}}:
        raise ValueError("冻结运行时间已变化")
    channels = json.loads((input_dir / "Slack_Channels.json").read_text(encoding="utf-8"))
    if channels != {"channels": [{"id": "C_HR_5117", "name": "hr-general", "is_private": False}], "messages": []}:
        raise ValueError("Slack 频道快照已变化")
    workbook = load_workbook(input_dir / "Grievance_Tracker.xlsx", read_only=True, data_only=True)
    if workbook.sheetnames != ["Pending Cases"]:
        raise ValueError("申诉台账工作表结构已变化")
    cases = [GrievanceCase(int(row[0]), *(str(value or "") for value in row[1:])) for row in workbook["Pending Cases"].iter_rows(min_row=2, values_only=True) if row[0] is not None]
    workbook.close()
    expected_cases = [
        (2,"GR-2026-041","Anna Svensson","Robert Kim","robert.kim@company.example.com","Workplace Conduct","Level 1","Noise complaint about adjacent team area","Pending",""),
        (3,"GR-2026-042","Ben Torres","Lisa Wang","lisa.wang@company.example.com","Compensation","Level 2","Overtime not being properly tracked","Pending",""),
        (4,"GR-2026-043","Carmen Obi","Derek Huang","derek.huang@company.example.com","Discrimination","Level 1","Feels excluded from team activities","Pending",""),
        (5,"GR-2026-044","David Holm","Mira Patel","mira.patel@company.example.com","Safety","Level 3","Exposed wiring in server room","Pending",""),
        (6,"GR-2026-045","Elise Braun","Robert Kim","robert.kim@company.example.com","Workplace Conduct","Level 1","Manager Robert Kim making inappropriate comments","Pending","Complaint is about direct manager"),
        (7,"GR-2026-046","Frank Yoshida","Lisa Wang","lisa.wang@company.example.com","Policy Dispute","Level 2","Disagrees with new remote work policy application","Pending",""),
        (8,"GR-2026-047","Grace Petrov","Derek Huang","derek.huang@company.example.com","Discrimination","Level 2","Passed over for promotion due to gender","Pending",""),
        (9,"GR-2026-048","Hugo Strand","Mira Patel","mira.patel@company.example.com","Workplace Conduct","Level 2","Persistent bullying from peer","Pending",""),
    ]
    if [tuple(item.__dict__.values()) for item in cases] != expected_cases:
        raise ValueError("八个冻结申诉案件已变化")
    box = mailbox.mbox(input_dir / "Grievance_Procedures.mbox", create=False)
    source_messages = list(box); box.close()
    if len(source_messages) != 2:
        raise ValueError("Grievance Procedures MBOX 必须恰好两封")
    messages = [{"message_id":str(message.get("X-AutomationBench-Message-ID","")),"thread_id":str(message.get("X-AutomationBench-Thread-ID","")),"from":str(message.get("From","")),"to":str(message.get("To","")),"subject":str(message.get("Subject","")),"date":str(message.get("Date","")),"body":decode_message(message)} for message in source_messages]
    if [item["message_id"] for item in messages] != ["msg_grievance_handbook_5117", "msg_griev_override_5117"]:
        raise ValueError("两封权威邮件身份已变化")
    if not all(token in messages[0]["body"] for token in ("Level 1", "Level 2", "Level 3", "strictly confidential", "Do not post", "manager IS the subject")):
        raise ValueError("申诉标准程序不完整")
    if not all(token in messages[1]["body"] for token in ("ALL discrimination-category grievances", "Legal AND HR Director", "overrides")):
        raise ValueError("歧视类覆盖规则不完整")
    return cases, messages, channels, context["meta"]["current_time"]


def derive_truth(cases: list[GrievanceCase]) -> list[Decision]:
    decisions: list[Decision] = []
    for item in cases:
        discrimination = item.category == "Discrimination"
        manager_subject = "direct manager" in item.notes.lower()
        if discrimination:
            route, handler, to_email, cc_email = "LEGAL_HR_DIRECTOR", "Legal and HR Director", "legal@company.example.com", "hr-director@company.example.com"
            authority = "HR Director 2026-04-01 discrimination override"
        elif manager_subject:
            route, handler, to_email, cc_email = "HRBP", "HRBP", "hr-bp@company.example.com", ""
            authority = "Handbook Section 6 manager-subject exception"
        elif item.severity == "Level 1":
            route, handler, to_email, cc_email = "MANAGER", item.manager, item.manager_email, ""
            authority = "Handbook Section 4 Level 1"
        elif item.severity == "Level 2":
            route, handler, to_email, cc_email = "HRBP", "HRBP", "hr-bp@company.example.com", ""
            authority = "Handbook Section 4 Level 2"
        elif item.severity == "Level 3":
            route, handler, to_email, cc_email = "LEGAL_HR_DIRECTOR", "Legal and HR Director", "legal@company.example.com", "hr-director@company.example.com"
            authority = "Handbook Section 4 Level 3"
        else:
            raise ValueError(f"未知严重度：{item.severity}")
        decisions.append(Decision(item,"YES" if discrimination else "NO","YES" if manager_subject else "NO",route,handler,to_email,cc_email,f"Routed to {handler}",authority))
    expected = [
        ("GR-2026-041","MANAGER","Robert Kim","robert.kim@company.example.com",""),
        ("GR-2026-042","HRBP","HRBP","hr-bp@company.example.com",""),
        ("GR-2026-043","LEGAL_HR_DIRECTOR","Legal and HR Director","legal@company.example.com","hr-director@company.example.com"),
        ("GR-2026-044","LEGAL_HR_DIRECTOR","Legal and HR Director","legal@company.example.com","hr-director@company.example.com"),
        ("GR-2026-045","HRBP","HRBP","hr-bp@company.example.com",""),
        ("GR-2026-046","HRBP","HRBP","hr-bp@company.example.com",""),
        ("GR-2026-047","LEGAL_HR_DIRECTOR","Legal and HR Director","legal@company.example.com","hr-director@company.example.com"),
        ("GR-2026-048","HRBP","HRBP","hr-bp@company.example.com",""),
    ]
    actual = [(item.source.case_id,item.route_type,item.handler,item.to_email,item.cc_email) for item in decisions]
    if actual != expected:
        raise ValueError("八案路由真值不一致")
    return decisions


def add_table(sheet, name: str) -> None:
    table = Table(displayName=name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)


def style_sheet(sheet, widths: dict[str, float], footer: str, *, a3: bool = False) -> None:
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions; sheet.sheet_view.showGridLines = False
    for column, width in widths.items(): sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY); cell.font = Font(name="Aptos", bold=True, color=WHITE, size=8); cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = Border(bottom=THIN)
    sheet.row_dimensions[1].height = 31
    grid = Side(style="hair", color="D9E2F3")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=7.8); cell.alignment = Alignment(vertical="top", wrap_text=True); cell.border = Border(left=grid, right=grid, bottom=grid)
    sheet.page_setup.orientation = "landscape"; sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if a3 else sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 1; sheet.sheet_properties.pageSetUpPr.fitToPage = True; sheet.print_options.horizontalCentered = True
    sheet.oddFooter.center.text = footer; sheet.print_area = sheet.dimensions


def normalize_ooxml(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="grievance-5117-ooxml-") as temporary:
        target = Path(temporary) / path.name; axis_ids: dict[str, str] = {}
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as destination:
            for name in sorted(source.namelist()):
                data = source.read(name)
                if name == "docProps/core.xml":
                    try:
                        root = etree.fromstring(data); ns = {"dcterms":"http://purl.org/dc/terms/","cp":"http://schemas.openxmlformats.org/package/2006/metadata/core-properties"}
                        for tag in ("created", "modified"):
                            for node in root.xpath(f"//dcterms:{tag}", namespaces=ns): node.text = FIXED_ISO_TIME
                        for tag in ("lastModifiedBy", "revision"):
                            for node in root.xpath(f"//cp:{tag}", namespaces=ns): node.text = "EvoCowork" if tag == "lastModifiedBy" else "1"
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except etree.XMLSyntaxError: pass
                if name.startswith("xl/charts/chart") and name.endswith(".xml"):
                    try:
                        root = etree.fromstring(data); ns = {"c":"http://schemas.openxmlformats.org/drawingml/2006/chart"}
                        for node in root.xpath("//c:axId", namespaces=ns):
                            old = node.get("val")
                            if old and old not in axis_ids: axis_ids[old] = str(51_170_001 + len(axis_ids))
                        for node in root.xpath("//c:axId | //c:crossAx", namespaces=ns):
                            old = node.get("val")
                            if old in axis_ids: node.set("val", axis_ids[old])
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except etree.XMLSyntaxError: pass
                info = zipfile.ZipInfo(name, FIXED_ZIP_TIME); info.compress_type = zipfile.ZIP_DEFLATED; info.external_attr = 0o600 << 16; destination.writestr(info, data)
        shutil.copy2(target, path)


def recalculate_xlsx(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="grievance-5117-recalc-") as temporary:
        root = Path(temporary); source_dir = root / "source"; output_dir = root / "output"; profile_dir = root / "profile"
        source_dir.mkdir(); output_dir.mkdir(); source = source_dir / path.name; shutil.copy2(path, source)
        process = subprocess.run(["libreoffice","--headless",f"-env:UserInstallation=file://{profile_dir}","--convert-to","xlsx","--outdir",str(output_dir),str(source)],stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,timeout=300)
        converted = output_dir / path.name
        if process.returncode != 0 or not converted.is_file(): raise RuntimeError(f"LibreOffice 公式重算失败：{process.stdout[-1000:]}")
        shutil.copy2(converted, path)


def write_workbook(path: Path, cases: list[GrievanceCase], messages: list[dict], channels: dict, frozen_time: str) -> None:
    workbook = Workbook(); workbook.remove(workbook.active)
    workbook.creator = "EvoCowork AutomationBench 5117"; workbook.title = "Grievance routing and confidentiality audit"; workbook.subject = "Eight-case deterministic routing under policy precedence"; workbook.description = "Restricted case truth, formula-derived routes, authorized draft queues and zero-side-effect controls."
    source = workbook.create_sheet("Restricted_Source"); source.append(["Source Row","Case ID","Employee","Manager","Manager Email","Category","Severity","Description","Status","Notes"])
    for item in cases: source.append(list(item.__dict__.values()))
    add_table(source,"RestrictedGrievanceSource"); style_sheet(source,{"A":10,"B":17,"C":19,"D":18,"E":31,"F":22,"G":12,"H":46,"I":12,"J":37},"RESTRICTED · AutomationBench 5117 source",a3=True)

    policy = workbook.create_sheet("Policy_Source"); policy.append(["Message ID","From","To","Subject","Date","Body","Source Trace"])
    for item in messages: policy.append([item["message_id"],item["from"],item["to"],item["subject"],item["date"],item["body"],f"Grievance_Procedures.mbox:{item['message_id']}"])
    add_table(policy,"GrievancePolicySource"); style_sheet(policy,{"A":31,"B":31,"C":31,"D":48,"E":28,"F":110,"G":50},"RESTRICTED · exact policy evidence",a3=True)
    for row in range(2,4): policy.row_dimensions[row].height = 165

    evidence = workbook.create_sheet("Evidence_Register"); evidence.append(["Evidence ID","Type","Source","Exact Fact","Used By"])
    for item in cases: evidence.append([f"EV-CASE-{item.source_row}","CASE","Grievance_Tracker.xlsx:Pending Cases",f"row={item.source_row}; case={item.case_id}; employee={item.employee}; manager={item.manager}; manager_email={item.manager_email}; category={item.category}; severity={item.severity}; description={item.description}; status={item.status}; notes={item.notes}","Routing_Decisions"])
    for index,item in enumerate(messages,1): evidence.append([f"EV-POLICY-{index}","POLICY",f"Grievance_Procedures.mbox:{item['message_id']}",item["body"],"Routing_Decisions / notifications / controls"])
    evidence.append(["EV-CONTEXT-1","TIME","Workflow_Context.json",f"current_time={frozen_time}","All deliverables"])
    channel = channels["channels"][0]; evidence.append(["EV-SLACK-1","CHANNEL","Slack_Channels.json",json.dumps(channel,ensure_ascii=False,separators=(",",":")),"Grievance_Slack_Control.json"])
    add_table(evidence,"GrievanceEvidenceRegister"); style_sheet(evidence,{"A":20,"B":14,"C":47,"D":118,"E":45},"RESTRICTED · source-backed evidence register",a3=True)
    for row in range(2,14): evidence.row_dimensions[row].height = 63 if row < 10 else 115

    decisions = workbook.create_sheet("Routing_Decisions"); decisions.append(["Source Row","Case ID","Employee","Manager","Manager Email","Category","Severity","Description","Source Status","Notes","Discrimination Override","Manager Subject","Route Type","Primary Handler","To","Cc","Proposed Status","Authority","Source Trace","Slack Allowed","Confidentiality"])
    for row in range(2,10):
        src = row
        values = [f"=Restricted_Source!A{src}",f"=Restricted_Source!B{src}",f"=Restricted_Source!C{src}",f"=Restricted_Source!D{src}",f"=Restricted_Source!E{src}",f"=Restricted_Source!F{src}",f"=Restricted_Source!G{src}",f"=Restricted_Source!H{src}",f"=Restricted_Source!I{src}",f'=IF(Restricted_Source!J{src}="","",Restricted_Source!J{src})',f'=IF(F{row}="Discrimination","YES","NO")',f'=IF(ISNUMBER(SEARCH("direct manager",J{row})),"YES","NO")',f'=IF(K{row}="YES","LEGAL_HR_DIRECTOR",IF(L{row}="YES","HRBP",IF(G{row}="Level 1","MANAGER",IF(G{row}="Level 2","HRBP",IF(G{row}="Level 3","LEGAL_HR_DIRECTOR","ERROR")))))',f'=IF(M{row}="MANAGER",D{row},IF(M{row}="HRBP","HRBP","Legal and HR Director"))',f'=IF(M{row}="MANAGER",E{row},IF(M{row}="HRBP","hr-bp@company.example.com","legal@company.example.com"))',f'=IF(M{row}="LEGAL_HR_DIRECTOR","hr-director@company.example.com","")',f'="Routed to "&N{row}',f'=IF(K{row}="YES","HR Director 2026-04-01 discrimination override",IF(L{row}="YES","Handbook Section 6 manager-subject exception",IF(G{row}="Level 1","Handbook Section 4 Level 1",IF(G{row}="Level 2","Handbook Section 4 Level 2","Handbook Section 4 Level 3"))))',f'="Grievance_Tracker.xlsx:Pending Cases row "&A{row}&" | Grievance_Procedures.mbox:"&IF(K{row}="YES","msg_griev_override_5117","msg_grievance_handbook_5117")',"NO","STRICTLY_CONFIDENTIAL"]
        decisions.append(values)
    add_table(decisions,"GrievanceRoutingDecisions"); style_sheet(decisions,{"A":9,"B":16,"C":19,"D":18,"E":29,"F":21,"G":11,"H":42,"I":11,"J":32,"K":16,"L":14,"M":22,"N":24,"O":32,"P":34,"Q":29,"R":47,"S":70,"T":12,"U":24},"RESTRICTED · formula-derived routing truth",a3=True)
    decisions.conditional_formatting.add("K2:K9",FormulaRule(formula=['K2="YES"'],fill=PatternFill("solid",fgColor=RED)))
    decisions.conditional_formatting.add("L2:L9",FormulaRule(formula=['L2="YES"'],fill=PatternFill("solid",fgColor=AMBER)))
    decisions.conditional_formatting.add("M2:M9",FormulaRule(formula=['M2="HRBP"'],fill=PatternFill("solid",fgColor=GREEN)))
    decisions.conditional_formatting.add("M2:M9",FormulaRule(formula=['M2="LEGAL_HR_DIRECTOR"'],fill=PatternFill("solid",fgColor=RED)))

    queue_specs = [
        ("Manager_Queue","ManagerRouteQueue",[2]),
        ("HRBP_Queue","HRBPRouteQueue",[3,6,7,9]),
        ("Legal_HR_Director_Queue","LegalHRDirectorRouteQueue",[4,5,8]),
    ]
    for title,table_name,rows in queue_specs:
        sheet = workbook.create_sheet(title); sheet.append(["Order","Source Row","Case ID","Route Type","Handler","To","Cc","Proposed Status","Authority","Draft State","Source Trace"])
        for order,decision_row in enumerate(rows,1): sheet.append([order,f"=Routing_Decisions!A{decision_row}",f"=Routing_Decisions!B{decision_row}",f"=Routing_Decisions!M{decision_row}",f"=Routing_Decisions!N{decision_row}",f"=Routing_Decisions!O{decision_row}",f"=Routing_Decisions!P{decision_row}",f"=Routing_Decisions!Q{decision_row}",f"=Routing_Decisions!R{decision_row}","DRAFT_NOT_SENT",f"=Routing_Decisions!S{decision_row}"])
        add_table(sheet,table_name); style_sheet(sheet,{"A":9,"B":11,"C":17,"D":23,"E":25,"F":33,"G":35,"H":31,"I":52,"J":21,"K":70},f"RESTRICTED · {title} · offline drafts",a3=True)

    summary = workbook.create_sheet("Leadership_Summary"); summary.append(["Route Type","Case Count","Case IDs","Policy Basis","Email Drafts","Emails Sent","Confidentiality"])
    summary_rows = [
        ("MANAGER",'=COUNTIF(Routing_Decisions!$M$2:$M$9,A2)','GR-2026-041','Handbook Section 4 Level 1',1,0,'AGGREGATE_CONTROL_VIEW'),
        ("HRBP",'=COUNTIF(Routing_Decisions!$M$2:$M$9,A3)','GR-2026-042, GR-2026-045, GR-2026-046, GR-2026-048','Handbook Section 4 Level 2 + Section 6 exception',4,0,'AGGREGATE_CONTROL_VIEW'),
        ("LEGAL_HR_DIRECTOR",'=COUNTIF(Routing_Decisions!$M$2:$M$9,A4)','GR-2026-043, GR-2026-044, GR-2026-047','Discrimination override + Handbook Section 4 Level 3',3,0,'AGGREGATE_CONTROL_VIEW'),
    ]
    for row in summary_rows: summary.append(row)
    add_table(summary,"GrievanceLeadershipSummary"); style_sheet(summary,{"A":25,"B":13,"C":64,"D":65,"E":14,"F":13,"G":28},"CONTROL VIEW · no employee or manager identities",a3=True)
    chart = BarChart(); chart.type="bar"; chart.style=10; chart.title="Cases by authorized route"; chart.y_axis.title="Route"; chart.x_axis.title="Cases"; chart.height=6.5; chart.width=12
    chart.add_data(Reference(summary,min_col=2,min_row=1,max_row=4),titles_from_data=True); chart.set_categories(Reference(summary,min_col=1,min_row=2,max_row=4)); summary.add_chart(chart,"I2")

    control = workbook.create_sheet("Control_Summary"); control.append(["Metric","Value","Expected","Status","Source Trace"])
    metrics = [
        ("Source cases",'=COUNTA(Restricted_Source!B2:B9)',8),("Level 1 cases",'=COUNTIF(Routing_Decisions!G2:G9,"Level 1")',3),("Level 2 cases",'=COUNTIF(Routing_Decisions!G2:G9,"Level 2")',4),("Level 3 cases",'=COUNTIF(Routing_Decisions!G2:G9,"Level 3")',1),
        ("Discrimination overrides",'=COUNTIF(Routing_Decisions!K2:K9,"YES")',2),("Manager-subject exceptions",'=COUNTIF(Routing_Decisions!L2:L9,"YES")',1),("Manager routes",'=COUNTIF(Routing_Decisions!M2:M9,"MANAGER")',1),("HRBP routes",'=COUNTIF(Routing_Decisions!M2:M9,"HRBP")',4),("Legal + HR Director routes",'=COUNTIF(Routing_Decisions!M2:M9,"LEGAL_HR_DIRECTOR")',3),
        ("Email drafts",'=COUNTA(Routing_Decisions!B2:B9)',8),("Emails sent",'=0',0),("Tracker update drafts",'=COUNTA(Routing_Decisions!Q2:Q9)',8),("Tracker updates applied",'=0',0),("Slack drafts",'=0',0),("Slack posts",'=0',0),("Source workbook modifications",'=0',0),("Real side effects",'=0',0),
        ("First case",'=Routing_Decisions!B2','GR-2026-041'),("Last case",'=Routing_Decisions!B9','GR-2026-048'),("Manager queue rows",'=COUNTA(Manager_Queue!C2:C2)',1),("HRBP queue rows",'=COUNTA(HRBP_Queue!C2:C5)',4),("Legal queue rows",'=COUNTA(Legal_HR_Director_Queue!C2:C4)',3),("Confidentiality",'="STRICTLY_CONFIDENTIAL"','STRICTLY_CONFIDENTIAL'),("Slack disclosed case details",'=0',0),("Policy messages",'=COUNTA(Policy_Source!A2:A3)',2),("Evidence rows",'=COUNTA(Evidence_Register!A2:A13)',12),
    ]
    for index,(metric,value,expected) in enumerate(metrics,2): control.append([metric,value,expected,f'=IF(B{index}=C{index},"PASS","FAIL")',"Routing_Decisions / queues / source evidence"])
    add_table(control,"GrievanceControlSummary"); style_sheet(control,{"A":36,"B":27,"C":27,"D":12,"E":55},"CONTROL VIEW · 26 source-backed controls",a3=True)
    control.conditional_formatting.add("D2:D27",FormulaRule(formula=['D2="FAIL"'],fill=PatternFill("solid",fgColor=RED)))

    workbook.calculation.fullCalcOnLoad = True; workbook.calculation.forceFullCalc = True; workbook.calculation.calcMode = "auto"
    workbook.save(path); recalculate_xlsx(path); normalize_ooxml(path)


def set_cell_shading(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr(); shading = properties.find(qn("w:shd"))
    if shading is None: shading = OxmlElement("w:shd"); properties.append(shading)
    shading.set(qn("w:fill"), color)


def add_doc_table(document: Document, headers: list[str], rows: list[tuple], widths: list[float] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers)); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.style = "Table Grid"
    for index,header in enumerate(headers):
        cell=table.rows[0].cells[index]; cell.text=header; set_cell_shading(cell,NAVY)
        for run in cell.paragraphs[0].runs: run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255); run.font.size=Pt(8)
    for values in rows:
        cells=table.add_row().cells
        for index,value in enumerate(values):
            cells[index].text=str(value); cells[index].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after=Pt(0)
                for run in paragraph.runs: run.font.name="Aptos"; run.font.size=Pt(7.4)
    if widths:
        for row in table.rows:
            for index,width in enumerate(widths): row.cells[index].width=Cm(width)


def set_doc_header_footer(document: Document) -> None:
    for section in document.sections:
        section.orientation=WD_ORIENT.LANDSCAPE; section.page_width=Cm(29.7); section.page_height=Cm(21); section.top_margin=Cm(1.25); section.bottom_margin=Cm(1.1); section.left_margin=Cm(1.25); section.right_margin=Cm(1.25)
        header=section.header.paragraphs[0]; header.text="EVOCOWORK · GRIEVANCE ROUTING CONTROL · RESTRICTED"; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs: run.font.name="Aptos"; run.font.size=Pt(7); run.font.color.rgb=RGBColor(91,111,131)
        footer=section.footer.paragraphs[0]; footer.text="AutomationBench 5117 · Offline draft only · 2026-04-15"; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for run in footer.runs: run.font.name="Aptos"; run.font.size=Pt(7); run.font.color.rgb=RGBColor(91,111,131)


def add_doc_title(document: Document, title: str, subtitle: str) -> None:
    paragraph=document.add_paragraph(); paragraph.paragraph_format.space_after=Pt(2); paragraph.alignment=WD_ALIGN_PARAGRAPH.LEFT
    run=paragraph.add_run(title); run.font.name="Aptos Display"; run.font.size=Pt(23); run.font.bold=True; run.font.color.rgb=RGBColor(23,54,93)
    paragraph=document.add_paragraph(subtitle); paragraph.paragraph_format.space_after=Pt(7)
    for run in paragraph.runs: run.font.name="Aptos"; run.font.size=Pt(9); run.font.color.rgb=RGBColor(80,96,112)


def write_brief(path: Path, decisions: list[Decision]) -> None:
    document=Document(); set_doc_header_footer(document)
    styles=document.styles
    styles["Normal"].font.name="Aptos"; styles["Normal"].font.size=Pt(8.2); styles["Normal"].paragraph_format.space_after=Pt(3)
    add_doc_title(document,"Employee grievance routing control","Eight source cases · policy precedence · confidential offline drafts")
    add_doc_table(document,["Route","Cases","Case IDs","Policy basis","Drafts","Sent"],[
        ("MANAGER",1,"GR-2026-041","Section 4 · Level 1",1,0),
        ("HRBP",4,"GR-2026-042, 045, 046, 048","Section 4 · Level 2 + Section 6",4,0),
        ("LEGAL_HR_DIRECTOR",3,"GR-2026-043, 044, 047","Discrimination override + Level 3",3,0),
    ],[3.8,2,7.6,8.3,2,2])
    document.add_paragraph("Control result: 8 of 8 cases have a single deterministic route; 2 discrimination overrides and 1 manager-subject exception are applied before standard severity routing. All notification and tracker actions remain offline drafts.")
    document.add_page_break(); add_doc_title(document,"Case routing matrix","Case-level control view without employee, manager, email or grievance-description disclosure")
    rows=[(d.source.case_id,d.source.category,d.source.severity,d.discrimination_override,d.manager_subject,d.route_type,d.authority) for d in decisions]
    add_doc_table(document,["Case ID","Category","Severity","Discrimination override","Manager subject","Route","Authority"],rows,[3.1,4.2,2.2,3.5,3.1,4.1,8.5])
    document.add_paragraph("Precedence: discrimination override -> manager-subject exception -> severity route. GR-2026-043 and GR-2026-047 therefore route to Legal + HR Director despite Level 1/2; GR-2026-045 skips the manager route and goes to HRBP.")
    document.add_page_break(); add_doc_title(document,"Confidentiality and execution boundary","No Slack grievance content · no source mutation · no external side effect")
    add_doc_table(document,["Control","Drafts","Executed","State"],[
        ("Notification emails",8,0,"DRAFT_NOT_SENT"),("Tracker updates",8,0,"DRAFT_NOT_APPLIED"),("Slack grievance messages",0,0,"BLOCKED_BY_POLICY"),("Source workbook changes",0,0,"UNCHANGED"),("Real side effects",0,0,"OFFLINE_ONLY")
    ],[8,3,3,8])
    add_doc_table(document,["Confidentiality rule","Applied result","Evidence"],[
        ("All grievance information is strictly confidential","Applied","Legal handbook Section 5"),("Do not post grievance details to any Slack channel","0 drafts · 0 posts","hr-general control record"),("Leadership brief excludes identities and descriptions","Applied","Case IDs + aggregate routes only"),("Authorized handler email contains case details","8 restricted drafts","One draft per source case")
    ],[9,6,10])
    document.add_paragraph("Final control: email drafts=8; emails sent=0; tracker update drafts=8; tracker updates applied=0; Slack drafts=0; Slack posts=0; source modifications=0; real side effects=0.")
    document.core_properties.title="Employee Grievance Routing Control Brief"; document.core_properties.subject="AutomationBench 5117 deterministic routing and confidentiality"; document.core_properties.author="EvoCowork"; document.core_properties.created=datetime(2026,4,15,9,0,0); document.core_properties.modified=datetime(2026,4,15,9,0,0)
    document.save(path); normalize_ooxml(path)


def add_textbox(slide,left,top,width,height,text,*,font_size=12,bold=False,color=PPT_NAVY,align=PP_ALIGN.LEFT,fill=None,line=None,radius=False):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(left),Inches(top),Inches(width),Inches(height)); shape.text=text; shape.text_frame.margin_left=Inches(0.10); shape.text_frame.margin_right=Inches(0.10); shape.text_frame.margin_top=Inches(0.06); shape.text_frame.margin_bottom=Inches(0.04); shape.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE
    for paragraph in shape.text_frame.paragraphs:
        paragraph.alignment=align
        for run in paragraph.runs: run.font.name="Aptos"; run.font.size=PptPt(font_size); run.font.bold=bold; run.font.color.rgb=color
    shape.fill.solid(); shape.fill.fore_color.rgb=fill or PPT_WHITE; shape.line.color.rgb=line or (fill or PPT_WHITE); return shape


def add_slide_title(slide,title,subtitle):
    add_textbox(slide,0.55,0.35,12.2,0.55,title,font_size=23,bold=True,color=PPT_NAVY)
    add_textbox(slide,0.55,0.9,12.2,0.34,subtitle,font_size=10,color=PptRGBColor(80,96,112))


def style_ppt_table(table,body_font=9):
    for row_index,row in enumerate(table.rows):
        for cell in row.cells:
            cell.margin_left=Inches(0.05); cell.margin_right=Inches(0.05); cell.margin_top=Inches(0.025); cell.margin_bottom=Inches(0.025); cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            cell.fill.solid(); cell.fill.fore_color.rgb=PPT_NAVY if row_index==0 else (PPT_LIGHT_BLUE if row_index%2 else PPT_WHITE)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment=PP_ALIGN.CENTER
                for run in paragraph.runs: run.font.name="Aptos"; run.font.size=PptPt(9 if row_index==0 else body_font); run.font.bold=row_index==0; run.font.color.rgb=PPT_WHITE if row_index==0 else PPT_NAVY


def write_presentation(path: Path, decisions: list[Decision]) -> None:
    presentation=Presentation(); presentation.slide_width=Inches(13.333); presentation.slide_height=Inches(7.5); blank=presentation.slide_layouts[6]
    slide=presentation.slides.add_slide(blank); add_slide_title(slide,"Employee grievance routing","8 cases · deterministic policy precedence · offline only")
    cards=[("8","source cases",PPT_LIGHT_BLUE),("1","manager",PPT_GREEN),("4","HRBP",PPT_AMBER),("3","Legal + HR Director",PPT_RED),("0","Slack posts",PPT_GREEN)]
    for index,(value,label,fill) in enumerate(cards):
        left=0.55+index*2.5; add_textbox(slide,left,1.35,2.15,0.75,value,font_size=24,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=fill,line=PPT_BLUE,radius=True); add_textbox(slide,left,2.1,2.15,0.33,label,font_size=9,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER)
    table=slide.shapes.add_table(4,6,Inches(0.55),Inches(2.65),Inches(12.2),Inches(2.5)).table
    rows=[("Route","Cases","Case IDs","Policy basis","Drafts","Sent"),("MANAGER",1,"GR-2026-041","Level 1",1,0),("HRBP",4,"042 · 045 · 046 · 048","Level 2 + Section 6",4,0),("LEGAL_HR_DIRECTOR",3,"043 · 044 · 047","Override + Level 3",3,0)]
    for r,values in enumerate(rows):
        for c,value in enumerate(values): table.cell(r,c).text=str(value)
    for index,width in enumerate([2.2,1.2,3.2,3.0,1.2,1.2]): table.columns[index].width=Inches(width)
    style_ppt_table(table,body_font=9); add_textbox(slide,0.55,5.55,12.2,0.7,"Precedence: discrimination override -> manager-subject exception -> severity route",font_size=13,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_LIGHT_BLUE,line=PPT_BLUE,radius=True)

    slide=presentation.slides.add_slide(blank); add_slide_title(slide,"Eight-case routing matrix","Case IDs and policy controls only · no identity or grievance-description disclosure")
    table=slide.shapes.add_table(9,6,Inches(0.55),Inches(1.35),Inches(12.2),Inches(4.8)).table
    headers=("Case ID","Category","Severity","Override","Manager subject","Route")
    for c,value in enumerate(headers): table.cell(0,c).text=value
    for r,d in enumerate(decisions,1):
        values=(d.source.case_id,d.source.category,d.source.severity,d.discrimination_override,d.manager_subject,d.route_type)
        for c,value in enumerate(values): table.cell(r,c).text=str(value)
    for index,width in enumerate([2.0,2.4,1.45,1.5,2.0,2.85]): table.columns[index].width=Inches(width)
    style_ppt_table(table,body_font=8.2); add_textbox(slide,0.55,6.35,12.2,0.45,"043 + 047 use the discrimination override · 045 skips its manager · 044 follows Level 3",font_size=10.5,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_AMBER,line=PPT_AMBER,radius=True)

    slide=presentation.slides.add_slide(blank); add_slide_title(slide,"Confidential delivery controls","Authorized email drafts only · Slack disclosure blocked · zero execution")
    table=slide.shapes.add_table(6,5,Inches(0.55),Inches(1.45),Inches(12.2),Inches(3.45)).table
    rows=[("Artifact","Drafts","Executed","Confidentiality","State"),("Notification emails",8,0,"AUTHORIZED_HANDLERS_ONLY","DRAFT_NOT_SENT"),("Tracker updates",8,0,"CASE_ID_CONTROL","DRAFT_NOT_APPLIED"),("Slack grievance messages",0,0,"NO_CASE_DETAILS","BLOCKED_BY_POLICY"),("Source workbook changes",0,0,"SOURCE_PRESERVED","UNCHANGED"),("Real side effects",0,0,"OFFLINE_ONLY","ZERO")]
    for r,values in enumerate(rows):
        for c,value in enumerate(values): table.cell(r,c).text=str(value)
    for index,width in enumerate([3.2,1.4,1.5,3.2,2.9]): table.columns[index].width=Inches(width)
    style_ppt_table(table,body_font=8.6)
    add_textbox(slide,0.55,5.3,3.8,0.8,"Email\n8 drafts · 0 sent",font_size=13,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_GREEN,line=PPT_BLUE,radius=True)
    add_textbox(slide,4.77,5.3,3.8,0.8,"Tracker\n8 drafts · 0 applied",font_size=13,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_AMBER,line=PPT_BLUE,radius=True)
    add_textbox(slide,8.99,5.3,3.8,0.8,"Slack\n0 drafts · 0 posts",font_size=13,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_GREEN,line=PPT_BLUE,radius=True)
    add_textbox(slide,0.55,6.38,12.2,0.42,"No employee or manager identities, personal email addresses, or grievance descriptions in this review",font_size=10.2,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_LIGHT_BLUE,line=PPT_LIGHT_BLUE,radius=True)
    presentation.core_properties.title="Employee Grievance Routing Review"; presentation.core_properties.subject="AutomationBench 5117 deterministic routing and confidentiality"; presentation.core_properties.author="EvoCowork"; presentation.core_properties.created=datetime(2026,4,15,9,0,0); presentation.core_properties.modified=datetime(2026,4,15,9,0,0)
    presentation.save(path); normalize_ooxml(path)


def email_body(decision: Decision) -> str:
    item=decision.source
    return "\n".join([
        "CONFIDENTIAL EMPLOYEE GRIEVANCE", "",
        f"Case ID: {item.case_id}", f"Employee: {item.employee}", f"Category: {item.category}", f"Severity: {item.severity}", f"Description: {item.description}",
        f"Routing authority: {decision.authority}", f"Proposed tracker status: {decision.proposed_status}", "",
        "Please handle this case under the confidential grievance procedure.", "DRAFT_NOT_SENT — this offline benchmark artifact has not been sent.", "- HR Operations",
    ])


def write_mbox(path: Path, decisions: list[Decision]) -> None:
    chunks=[]
    for order,decision in enumerate(decisions,1):
        case=decision.source; message=EmailMessage(policy=SMTP); message["From"]="hr-ops@company.example.com"; message["To"]=decision.to_email
        if decision.cc_email: message["Cc"]=decision.cc_email
        message["Subject"]=f"CONFIDENTIAL: Grievance {case.case_id} Routing Notification"; message["Date"]="Wed, 15 Apr 2026 09:00:00 +0000"; message["Message-ID"]=f"<grievance-5117-{case.case_id[-3:]}@evocowork.local>"; message["X-EvoCowork-Status"]="DRAFT_NOT_SENT"; message["X-EvoCowork-Route"]=decision.route_type; message["X-EvoCowork-Order"]=str(order); message["X-EvoCowork-Confidentiality"]="STRICTLY_CONFIDENTIAL"
        message.set_content(email_body(decision),charset="utf-8",cte="quoted-printable"); stream=BytesIO(); BytesGenerator(stream,policy=SMTP.clone(linesep="\n",max_line_length=78)).flatten(message,unixfrom=False)
        chunks.append(b"From MAILER-DAEMON Wed Apr 15 09:00:01 2026\n"+stream.getvalue().replace(b"\nFrom ",b"\n>From ")+b"\n")
    path.write_bytes(b"".join(chunks))


def write_updates(path: Path, decisions: list[Decision]) -> None:
    updates=[{"order":index,"source_row":item.source.source_row,"case_id":item.source.case_id,"route_type":item.route_type,"proposed_status":item.proposed_status,"authority":item.authority,"state":"DRAFT_NOT_APPLIED","applied":False,"source_trace":f"Grievance_Tracker.xlsx:Pending Cases row {item.source.source_row}"} for index,item in enumerate(decisions,1)]
    payload={"task_id":"automationbench-5117","status":"DRAFT_NOT_APPLIED","spreadsheet_id":"ss_grievances_5117","worksheet_id":"ws_cases_5117","update_count":8,"applied_count":0,"source_modified":False,"updates":updates,"source_trace":["Grievance_Routing_Audit.xlsx:Routing_Decisions","Grievance_Procedures.mbox:msg_grievance_handbook_5117","Grievance_Procedures.mbox:msg_griev_override_5117"]}
    path.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")


def write_slack(path: Path) -> None:
    payload={"task_id":"automationbench-5117","source_channel":{"id":"C_HR_5117","name":"hr-general","is_private":False},"status":"BLOCKED_BY_CONFIDENTIALITY_POLICY","draft_count":0,"posted_count":0,"posted":False,"case_detail_count":0,"employee_name_count":0,"manager_name_count":0,"messages":[],"reason":"Employee grievance procedures prohibit posting grievance details to any Slack channel.","controls":["NO_CASE_IDS","NO_EMPLOYEE_NAMES","NO_MANAGER_NAMES","NO_GRIEVANCE_DESCRIPTIONS","NO_SLACK_DRAFTS","NO_SLACK_POSTS"],"source_trace":["Grievance_Procedures.mbox:msg_grievance_handbook_5117 Section 5","Slack_Channels.json:C_HR_5117"]}
    path.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")


def write_execution(path: Path) -> None:
    payload={"task_id":"automationbench-5117","status":"CONTROLLED_DRAFT_ONLY","notification_email":{"draft_count":8,"sent_count":0,"state":"DRAFT_NOT_SENT"},"tracker_updates":{"draft_count":8,"applied_count":0,"state":"DRAFT_NOT_APPLIED"},"slack":{"channel":"C_HR_5117 / hr-general","draft_count":0,"posted_count":0,"state":"BLOCKED_BY_POLICY"},"board_artifacts":{"docx_count":1,"pptx_count":1,"state":"OFFLINE_ARTIFACTS"},"source_workbook":{"modified_row_count":0,"state":"UNCHANGED"},"confidentiality":{"authorized_email_case_count":8,"slack_case_detail_count":0,"leadership_identity_count":0,"state":"PRESERVED"},"real_side_effect_count":0,"source_trace":["Grievance_Tracker.xlsx:Pending Cases rows 2-9","Grievance_Procedures.mbox:msg_grievance_handbook_5117","Grievance_Procedures.mbox:msg_griev_override_5117","Slack_Channels.json:C_HR_5117"]}
    path.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")


def main() -> None:
    parser=argparse.ArgumentParser(); parser.add_argument("--input-dir",type=Path,required=True); parser.add_argument("--output-dir",type=Path,required=True); args=parser.parse_args()
    cases,messages,channels,frozen_time=read_inputs(args.input_dir); decisions=derive_truth(cases)
    args.output_dir.mkdir(parents=True,exist_ok=True)
    for item in args.output_dir.iterdir():
        if item.is_file(): item.unlink()
        elif item.is_dir(): shutil.rmtree(item)
    write_workbook(args.output_dir/XLSX_NAME,cases,messages,channels,frozen_time); write_brief(args.output_dir/DOCX_NAME,decisions); write_presentation(args.output_dir/PPTX_NAME,decisions); write_mbox(args.output_dir/MBOX_NAME,decisions); write_updates(args.output_dir/UPDATES_NAME,decisions); write_slack(args.output_dir/SLACK_NAME); write_execution(args.output_dir/EXECUTION_NAME)
    expected={XLSX_NAME,DOCX_NAME,PPTX_NAME,MBOX_NAME,UPDATES_NAME,SLACK_NAME,EXECUTION_NAME}
    if {item.name for item in args.output_dir.iterdir() if item.is_file()} != expected: raise RuntimeError("交付文件集合与冻结任务不一致")


if __name__ == "__main__": main()
