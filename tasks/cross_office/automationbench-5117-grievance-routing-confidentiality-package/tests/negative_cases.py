#!/usr/bin/env python3
"""AutomationBench 5117 的任务专属 v2 分层负例。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from typing import Callable

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
INPUTS = ROOT / "environment" / "input_files"
OUTPUTS = (
    "Grievance_Routing_Audit.xlsx",
    "Grievance_Routing_Control_Brief.docx",
    "Grievance_Routing_Review.pptx",
    "Grievance_Notification_Drafts.mbox",
    "Grievance_Tracker_Update_Drafts.json",
    "Grievance_Slack_Control.json",
    "Grievance_Execution_Control.json",
)
Mutator = Callable[[Path], None]


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"无法载入 {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _xlsx(root: Path, sheet: str, cell: str, value) -> None:
    path = root / OUTPUTS[0]
    book = load_workbook(path)
    book[sheet][cell] = value
    book.save(path)


def _json(root: Path, name: str, mutator: Callable[[dict], None]) -> None:
    path = root / name
    payload = json.loads(path.read_text(encoding="utf-8"))
    mutator(payload)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _rewrite_mbox(root: Path, mutator: Callable[[list], None]) -> None:
    path = root / OUTPUTS[3]
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    mutator(messages)
    with path.open("wb") as handle:
        for index, message in enumerate(messages, 1):
            handle.write(f"From MAILER-DAEMON Wed Apr 15 09:00:{index:02d} 2026\n".encode())
            BytesGenerator(handle, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
            handle.write(b"\n")


def _replace_word(root: Path, old: str, new: str) -> None:
    path = root / OUTPUTS[1]
    document = Document(path)
    changed = False
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if old in cell.text:
                    cell.text = cell.text.replace(old, new, 1)
                    changed = True
    if not changed:
        raise ValueError(f"Word中找不到{old!r}")
    document.save(path)


def _replace_ppt(root: Path, old: str, new: str) -> None:
    path = root / OUTPUTS[2]
    deck = Presentation(path)
    changed = False
    for slide in deck.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and old in shape.text:
                shape.text = shape.text.replace(old, new, 1)
                changed = True
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        if old in cell.text:
                            cell.text = cell.text.replace(old, new, 1)
                            changed = True
    if not changed:
        raise ValueError(f"PPT中找不到{old!r}")
    deck.save(path)


def _input_only(root: Path) -> None:
    for path in list(root.iterdir()):
        if path.is_file():
            path.unlink()
    for source in INPUTS.iterdir():
        if source.is_file():
            shutil.copy2(source, root / source.name)


def _cases() -> list[tuple[str, Mutator, str]]:
    cases: list[tuple[str, Mutator, str]] = []
    for name in OUTPUTS:
        cases.append((f"missing-{name}", lambda root, n=name: (root / n).unlink(), "fatal_zero"))
        cases.append((f"corrupt-{name}", lambda root, n=name: (root / n).write_bytes(b"not-a-valid-artifact"), "fatal_zero"))
    cases.extend([
        ("mbox-empty", lambda root: (root / OUTPUTS[3]).write_bytes(b""), "fatal_zero"),
        ("unmodified-input-only", _input_only, "fatal_zero"),
        ("source-041-category-local", lambda root: _xlsx(root, "Restricted_Source", "F2", "Policy Dispute"), "partial"),
        ("route-formula-041-static", lambda root: _xlsx(root, "Routing_Decisions", "M2", "MANAGER"), "partial"),
        ("route-043-override-lost", lambda root: _xlsx(root, "Routing_Decisions", "M4", "MANAGER"), "partial"),
        ("route-045-manager-conflict", lambda root: _xlsx(root, "Routing_Decisions", "M6", "MANAGER"), "partial"),
        ("mail-041-body-description-local", lambda root: _rewrite_mbox(root, lambda ms: ms[0].set_payload(ms[0].get_payload().replace("Noise complaint about adjacent team area", "Noise complaint"))), "partial"),
        ("mail-043-wrong-recipient", lambda root: _rewrite_mbox(root, lambda ms: ms[2].replace_header("To", "hr-bp@company.example.com")), "partial"),
        ("mail-047-missing-cc", lambda root: _rewrite_mbox(root, lambda ms: del_header(ms[6], "Cc")), "partial"),
        ("update-046-status-local", lambda root: _json(root, OUTPUTS[4], lambda p: p["updates"][5].update({"proposed_status": "Pending"})), "partial"),
        ("slack-case-leak", lambda root: _json(root, OUTPUTS[5], lambda p: p.update({"draft_count": 1, "messages": [{"text": "GR-2026-041"}]})), "partial"),
        ("execution-email-sent", lambda root: _json(root, OUTPUTS[6], lambda p: p["notification_email"].update({"sent_count": 1, "state": "SENT"})), "partial"),
        ("word-043-route-local", lambda root: _replace_word(root, "HR Director 2026-04-01 discrimination override", "Handbook Section 4 Level 1"), "partial"),
        ("ppt-047-route-local", lambda root: _replace_ppt(root, "GR-2026-047", "GR-2026-049"), "partial"),
        ("ppt-zero-execution-conflict", lambda root: _replace_ppt(root, "8 drafts · 0 sent", "8 drafts · 1 sent"), "partial"),
        ("updates-applied-severe", lambda root: _json(root, OUTPUTS[4], lambda p: p.update({"applied_count": 1, "source_modified": True})), "partial"),
        ("excel-leadership-count-conflict", lambda root: _xlsx(root, "Leadership_Summary", "B3", 3), "partial"),
        ("extra-file", lambda root: (root / "unexpected.txt").write_text("unexpected", encoding="utf-8"), "partial"),
    ])
    return cases


def del_header(message, name: str) -> None:
    if message.get(name) is not None:
        del message[name]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    judge = _load("grievance_5117_v2_judge", ROOT / "tests" / "judge.py")
    judge.V1.render = lambda path: ({".xlsx": 9, ".docx": 3, ".pptx": 3}[path.suffix.casefold()], "render-ok")
    if judge.evaluate(GOLDEN).get("reward") != 1.0:
        raise AssertionError("Golden v2必须为1.0")
    results: list[dict] = []
    for name, mutator, layer in _cases():
        with tempfile.TemporaryDirectory(prefix=f"grievance-5117-v2-{name}-") as temporary:
            output = Path(temporary) / "output"
            shutil.copytree(GOLDEN, output)
            mutator(output)
            result = judge.evaluate(output)
            reward = float(result.get("reward", 0))
            if layer == "fatal_zero" and reward != 0:
                raise AssertionError(f"{name}应硬归零，实际{reward}")
            if layer == "partial" and not (0 < reward < 1):
                raise AssertionError(f"{name}应保留连续部分分，实际{reward}")
            results.append({"case": name, "layer": layer, "reward": reward, "passed": bool(result.get("passed")), "penalties": result.get("penalties", [])})
    partial = [item["reward"] for item in results if item["layer"] == "partial"]
    report = {
        "total": len(results), "fatal_zero": sum(item["layer"] == "fatal_zero" for item in results),
        "partial_continuous": len(partial), "partial_reward_range": [min(partial), max(partial)], "cases": results,
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
