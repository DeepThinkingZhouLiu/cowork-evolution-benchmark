#!/usr/bin/env python3
"""只把 AutomationBench 5118 的公开初始状态静态化为本题输入。

本脚本只接受 ``hr.visa_expiration_monitoring``，逐项断言八条签证记录、
四条台账政策、移民律师的最新处理周期邮件、冻结时间、Slack 快照和九条
上游 assertions。它不是通用制题器，不会创造员工、日期、规则或动作。
"""

from __future__ import annotations

import hashlib
import importlib.util
import json
import shutil
import sys
import types
import zipfile
from datetime import datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
UPSTREAM = Path(__file__).with_name("upstream-hr-tasks.py")
SNAPSHOT = Path(__file__).with_name("upstream-task-5118.json")
LICENSE = Path(__file__).with_name("LICENSE-AutomationBench-MIT")
DATA_DIR = ROOT / "data" / "input_files"
ENV_DIR = ROOT / "environment" / "input_files"
FIXED_ZIP_TIME = (2026, 4, 15, 9, 0, 0)
SOURCE_SHA256 = "c173a36a03f43096bd122577584aa32f9df379d8760aa46c43be04c85d101779"
LICENSE_SHA256 = "c6f9d87aff6d627d5cd78cb09364dafed7d6f1c4bb7a91183d57377157e96ba4"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_task() -> dict:
    if sha256(UPSTREAM) != SOURCE_SHA256 or sha256(LICENSE) != LICENSE_SHA256:
        raise ValueError("冻结的 AutomationBench 源文件或 LICENSE 哈希错误")
    datasets = types.ModuleType("datasets")
    datasets.Dataset = object
    sys.modules["datasets"] = datasets
    for name in ("automationbench", "automationbench.domains", "automationbench.domains.hr"):
        sys.modules[name] = types.ModuleType(name)
    noise = types.ModuleType("automationbench.domains.hr._noise")
    noise.apply_noise = lambda value: value
    sys.modules[noise.__name__] = noise
    spec = importlib.util.spec_from_file_location("automationbench_hr_tasks_frozen_5118", UPSTREAM)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入冻结的 AutomationBench HR 源文件")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    task = module.get_task_5118()
    if task.get("example_id") != 5118 or task.get("task") != "hr.visa_expiration_monitoring":
        raise ValueError("冻结函数不再是 AutomationBench 5118")
    return task


def assert_shape(task: dict) -> None:
    prompt = task["prompt"][1]["content"]
    assert all(token in prompt for token in ("monthly visa expiration check", "sponsored employees", "appropriate action", "immigration compliance process"))
    state = task["info"]["initial_state"]
    assert state["meta"] == {"schema_version": "0.1.0", "current_time": "2026-04-15T09:00:00Z"}
    messages = state["gmail"]["messages"]
    assert [item["id"] for item in messages] == ["msg_immigration_update_5118"]
    assert messages[0]["from_"] == "immigration-attorney@lawfirm.example.com"
    assert all(token in messages[0]["body_plain"] for token in ("H-1B renewals", "8 months", "previously 6 months", "adjust"))
    spreadsheet = state["google_sheets"]["spreadsheets"][0]
    assert spreadsheet["id"] == "ss_visa_5118" and spreadsheet["title"] == "Visa Tracking"
    active, policy = spreadsheet["worksheets"]
    assert active["id"] == "ws_visas_5118" and active["title"] == "Active Visas"
    assert [item["row_id"] for item in active["rows"]] == list(range(2, 10))
    assert [item["cells"]["Employee"] for item in active["rows"]] == ["Akiko Yamamoto", "Boris Kuznetsov", "Chen Wei", "Deepa Krishnan", "Esteban Ruiz", "Fatima Hassan", "Gustavo Lima", "Haruto Tanaka"]
    assert [item["cells"]["Visa Type"] for item in active["rows"]] == ["H-1B", "H-1B", "L-1", "H-1B", "TN", "H-1B", "H-1B", "O-1"]
    assert [item["cells"]["Expiration Date"] for item in active["rows"]] == ["2026-12-15", "2027-03-01", "2026-06-30", "2026-05-30", "2026-09-15", "2027-06-01", "2026-04-20", "2028-01-15"]
    assert active["rows"][3]["cells"]["Notes"] == "Renewal already filed"
    assert policy["id"] == "ws_policy_5118" and policy["title"] == "Immigration Compliance Policy"
    assert [item["row_id"] for item in policy["rows"]] == [2, 3, 4, 5]
    policy_text = "\n".join(item["cells"]["Rule"] + "\n" + item["cells"]["Action"] for item in policy["rows"])
    assert all(token in policy_text for token in ("within 30 days", "legal@company.example.com", "Renewal Initiated", "H-1B processing takes 6 months", "Already Filed", "no further action"))
    assert state["slack"] == {"channels": [{"id": "C_HR_5118", "name": "hr-general", "is_private": False}], "messages": []}
    assertions = task["info"]["assertions"]
    assert len(assertions) == 9
    assert not any("Fatima Hassan" in json.dumps(item) for item in assertions)


def normalize_xlsx(path: Path) -> None:
    temporary = path.with_suffix(".normalized.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    temporary.replace(path)


def style_sheet(sheet, widths: dict[str, float]) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = sheet.dimensions


def write_tracker(state: dict, target: Path) -> None:
    source = state["google_sheets"]["spreadsheets"][0]
    workbook = Workbook()
    workbook.creator = "AutomationBench 5118 public fixture"
    workbook.title = source["title"]
    workbook.subject = "Static Office representation of the visa tracker and its policy"
    workbook.description = "No employee, visa, date, status, note or policy was added beyond AutomationBench 5118."
    active, policy = source["worksheets"]
    sheet = workbook.active
    sheet.title = active["title"]
    sheet.append(["Source Row", "Employee", "Employee Email", "Visa Type", "Expiration Date", "Status", "Notes"])
    for item in active["rows"]:
        cells = item["cells"]
        sheet.append([item["row_id"], cells["Employee"], cells["Employee Email"], cells["Visa Type"], cells["Expiration Date"], cells["Status"], cells["Notes"]])
    table = Table(displayName="VisaTrackerSourceInput", ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    style_sheet(sheet, {"A": 12, "B": 25, "C": 39, "D": 15, "E": 18, "F": 15, "G": 30})
    rules = workbook.create_sheet(policy["title"])
    rules.append(["Source Row", "Rule", "Action"])
    for item in policy["rows"]:
        rules.append([item["row_id"], item["cells"]["Rule"], item["cells"]["Action"]])
    table = Table(displayName="ImmigrationPolicySourceInput", ref=rules.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    rules.add_table(table)
    style_sheet(rules, {"A": 12, "B": 53, "C": 120})
    for row in range(2, 6):
        rules.row_dimensions[row].height = 76
    workbook.save(target)
    normalize_xlsx(target)


def message_bytes(record: dict) -> bytes:
    message = EmailMessage(policy=SMTP)
    message["From"] = record["from_"]
    message["To"] = ", ".join(record["to"])
    message["Subject"] = record["subject"]
    parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
    message["Date"] = parsed.strftime("%a, %d %b %Y %H:%M:%S +0000")
    message["Message-ID"] = f"<{record['id']}@automationbench.invalid>"
    message["X-AutomationBench-Message-ID"] = record["id"]
    message["X-AutomationBench-Thread-ID"] = record["thread_id"]
    message["X-AutomationBench-Labels"] = ",".join(record.get("label_ids", []))
    message.set_content(record["body_plain"], charset="utf-8")
    buffer = BytesIO()
    BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
    return buffer.getvalue()


def write_mbox(messages: list[dict], target: Path) -> None:
    chunks: list[bytes] = []
    for record in messages:
        parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
        chunks.append(f"From {record['from_']} {parsed.strftime('%a %b %d %H:%M:%S %Y')}\n".encode("ascii"))
        chunks.append(message_bytes(record).replace(b"\r\n", b"\n"))
        chunks.append(b"\n")
    target.write_bytes(b"".join(chunks))


def main() -> None:
    task = load_task()
    assert_shape(task)
    state = task["info"]["initial_state"]
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    ENV_DIR.mkdir(parents=True, exist_ok=True)
    SNAPSHOT.write_text(json.dumps(task, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_tracker(state, DATA_DIR / "Visa_Tracking.xlsx")
    write_mbox(state["gmail"]["messages"], DATA_DIR / "Immigration_Update.mbox")
    (DATA_DIR / "Slack_Channels.json").write_text(json.dumps(state["slack"], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (DATA_DIR / "Workflow_Context.json").write_text(json.dumps({"meta": state["meta"]}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    for source in sorted(DATA_DIR.iterdir()):
        if source.is_file():
            shutil.copy2(source, ENV_DIR / source.name)


if __name__ == "__main__":
    main()
