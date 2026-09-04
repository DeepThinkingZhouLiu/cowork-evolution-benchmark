#!/usr/bin/env python3
"""AutomationBench 5118 的 v2 连续评分入口；原严格口径保存在 judge_v1.py。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import mailbox
import re
import sys
import zipfile
from email.utils import getaddresses
from pathlib import Path
from typing import Any, Callable

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

TASK_ID = "cowork-evo/automationbench-5118-visa-expiration-control-package"
RUBRIC_VERSION = "automationbench-5118-visa-expiration-v0.2"
JUDGE_VERSION = "automationbench-5118-criterion-partial-v2.0"
XLSX = "Visa_Expiration_Audit.xlsx"
DOCX = "Visa_Expiration_Control_Brief.docx"
PPTX = "Visa_Expiration_Review.pptx"
MBOX = "Visa_Notification_Drafts.mbox"
UPDATES = "Visa_Tracker_Update_Drafts.json"
SLACK = "Visa_Slack_Control.json"
EXECUTION = "Visa_Execution_Control.json"
REQUIRED = (XLSX, DOCX, PPTX, MBOX, UPDATES, SLACK, EXECUTION)
POSITIVE_WEIGHTS = {
    "R001": 5, "R002": 8, "R003": 11, "R004": 17,
    "R005": 12, "R006": 9, "R007": 8, "R008": 7,
    "R009": 6, "R010": 7, "R011": 4, "R012": 6,
}
PENALTY_WEIGHTS = {"P001": -35, "P002": -35, "P003": -10}


def _load_v1():
    path = Path(__file__).with_name("judge_v1.py")
    spec = importlib.util.spec_from_file_location("automationbench_5118_judge_v1", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v1 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


V1 = _load_v1()
SOURCE = V1.SOURCE
POLICIES = V1.POLICIES
TRUTH = V1.TRUTH
ACTION_NAMES = V1.ACTION_NAMES


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _norm(value: Any) -> Any:
    return V1.norm(value)


def _rows(sheet, start: int = 2, end: int | None = None, columns: int | None = None) -> list[tuple[Any, ...]]:
    end = end or sheet.max_row
    columns = columns or sheet.max_column
    return [
        tuple(_norm(sheet.cell(row, column).value) for column in range(1, columns + 1))
        for row in range(start, end + 1)
    ]


def _formula(value: Any, *tokens: str) -> bool:
    text = str(value or "")
    return text.startswith("=") and all(token.casefold() in text.casefold() for token in tokens)


def _messages(path: Path) -> list:
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    return messages


def _addresses(message, header: str) -> set[str]:
    return {address.casefold() for _, address in getaddresses(message.get_all(header, [])) if address}


def _criterion(key: str, score: float, evidence: str, earned: float = 0, total: float = 1) -> dict[str, Any]:
    return {
        "criterion_id": key, "score": round(max(0.0, min(1.0, score)), 6),
        "evidence": evidence, "method": "deterministic",
        "raw": {"earned_units": round(earned, 6), "total_units": round(total, 6)},
    }


def _fatal(errors: list[str], hashes: list[str]) -> dict[str, Any]:
    return {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes or ["0" * 64],
        "criterion_results": [_criterion(key, 0, "；".join(errors), 0, 0) for key in (*POSITIVE_WEIGHTS, *PENALTY_WEIGHTS)],
        "reward": 0.0, "passed": False, "hurdle_failures": ["R001"], "penalties": [],
    }


def _core(artifact: Path) -> tuple[bool, list[str]]:
    errors: list[str] = []
    if not artifact.is_dir():
        return False, ["输出目录不存在"]
    for name in REQUIRED:
        path = artifact / name
        if not path.is_file() or path.stat().st_size == 0:
            errors.append(f"核心交付缺失或为空：{name}")
            continue
        try:
            if name == XLSX:
                load_workbook(path, read_only=True, data_only=False).close()
            elif name == DOCX:
                Document(path)
            elif name == PPTX:
                if not Presentation(path).slides:
                    raise ValueError("PPT无页面")
            elif name == MBOX:
                messages = _messages(path)
                if not messages or not all(message.get("To") and message.get("Subject") for message in messages):
                    raise ValueError("MBOX无有效草稿")
            else:
                payload = json.loads(path.read_text(encoding="utf-8"))
                if not isinstance(payload, dict):
                    raise ValueError("JSON顶层不是对象")
        except Exception as exc:
            errors.append(f"核心交付无法解析：{name}: {exc}")
    return not errors, errors


def _ratio(findings: list[tuple[bool, str]]) -> tuple[float, list[str], int, int]:
    earned = sum(bool(ok) for ok, _ in findings)
    return earned / len(findings), [label for ok, label in findings if not ok], earned, len(findings)


def _leaf_findings(observed: Any, expected: Any, path: str) -> list[tuple[bool, str]]:
    if isinstance(expected, dict):
        findings = [(isinstance(observed, dict) and set(observed) == set(expected), f"{path}字段集合错误")]
        source = observed if isinstance(observed, dict) else {}
        for key, value in expected.items():
            findings.extend(_leaf_findings(source.get(key), value, f"{path}.{key}"))
        return findings
    if isinstance(expected, list):
        findings = [(isinstance(observed, list) and len(observed) == len(expected), f"{path}列表长度错误")]
        source = observed if isinstance(observed, list) else []
        for index, value in enumerate(expected):
            findings.extend(_leaf_findings(source[index] if index < len(source) else None, value, f"{path}[{index}]"))
        return findings
    return [(observed == expected, f"{path}应为{expected!r}")]


def _expected_policy() -> list[tuple[Any, ...]]:
    return [
        (row_id, rule, action, f"Visa_Tracking.xlsx:Immigration Compliance Policy row {row_id}")
        for row_id, rule, action in POLICIES
    ]


def _expected_evidence() -> list[tuple[Any, ...]]:
    update = V1.source_update()
    result = []
    for source_row, employee, email, visa, expiration, status, notes in SOURCE:
        result.append((
            f"EV-ROW-{source_row}", "VISA_RECORD", employee,
            f"employee={employee}; email={email}; visa={visa}; expiration={expiration}; status={status}; notes={notes}",
            f"Visa_Tracking.xlsx:Active Visas row {source_row}",
        ))
    for row_id, rule, action in POLICIES:
        result.append((f"EV-POLICY-{row_id}", "POLICY", rule, action, f"Visa_Tracking.xlsx:Immigration Compliance Policy row {row_id}"))
    result.extend([
        ("EV-ATTORNEY-1", "OVERRIDE", update["subject"], update["body"], "Immigration_Update.mbox:msg_immigration_update_5118"),
        ("EV-TIME-1", "TIME", "Frozen runtime", "2026-04-15T09:00:00Z", "Workflow_Context.json:meta.current_time"),
        ("EV-SLACK-1", "CHANNEL", "C_HR_5118 / hr-general", '{"id":"C_HR_5118","name":"hr-general","is_private":false}', "Slack_Channels.json:C_HR_5118"),
    ])
    return result


def _source_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    book = load_workbook(artifact / XLSX, data_only=True)
    source_rows = _rows(book["Source_Visas"], 2, 9, 7)
    policy_rows = _rows(book["Source_Policy"], 2, 5, 4)
    attorney_rows = _rows(book["Attorney_Update_Source"], 2, 2, 7)
    evidence_rows = _rows(book["Evidence_Register"], 2, 16, 5)
    book.close()
    findings: list[tuple[bool, str]] = []
    for index, wanted in enumerate(SOURCE):
        observed = source_rows[index] if index < len(source_rows) else ()
        for column, value in enumerate(wanted):
            findings.append((column < len(observed) and observed[column] == value, f"{wanted[1]}来源字段{column + 1}错误"))
    for index, wanted in enumerate(_expected_policy()):
        observed = policy_rows[index] if index < len(policy_rows) else ()
        for column, value in enumerate(wanted):
            findings.append((column < len(observed) and observed[column] == value, f"政策行{wanted[0]}字段{column + 1}错误"))
    update = V1.source_update()
    wanted_attorney = (update["message_id"], update["from"], update["to"], update["subject"], update["date"], update["body"], "Immigration_Update.mbox:msg_immigration_update_5118")
    observed = attorney_rows[0] if attorney_rows else ()
    for column, value in enumerate(wanted_attorney):
        findings.append((column < len(observed) and observed[column] == value, f"律师更新字段{column + 1}错误"))
    wanted_evidence = _expected_evidence()
    findings.append((len(evidence_rows) == 15, "Evidence_Register不是15条"))
    for index, wanted in enumerate(wanted_evidence):
        observed = evidence_rows[index] if index < len(evidence_rows) else ()
        for column, value in enumerate(wanted):
            findings.append((column < len(observed) and observed[column] == value, f"证据{wanted[0]}字段{column + 1}错误"))
    for filename, digest in V1.INPUT_HASHES.items():
        path = V1.input_root() / filename
        findings.append((path.is_file() and _sha256(path) == digest, f"冻结输入{filename}哈希错误"))
    return _ratio(findings)


def _formula_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    book = load_workbook(artifact / XLSX, data_only=False)
    names = ["Source_Visas", "Source_Policy", "Attorney_Update_Source", "Evidence_Register", "Visa_Decisions", "Urgent_Queue", "Renewal_Queue", "Already_Filed_Control", "Monitoring_Queue", "Control_Summary"]
    findings: list[tuple[bool, str]] = [(book.sheetnames == names, "Excel十表名称或顺序错误")]
    sheet = book["Visa_Decisions"]
    formula_rows = _rows(sheet, 2, 9, 25)
    for row_index, row in enumerate(formula_rows, 2):
        for column, value in enumerate(row, 1):
            findings.append((str(value).startswith("="), f"Visa_Decisions!{sheet.cell(row_index,column).coordinate}被静态化"))
    key_tokens = {
        5: ("DATEVALUE", "Source_Visas"), 8: ("DATE(2026,4,15)",), 9: ("E", "H"),
        10: ("H-1B", "L-1", "TN", "O-1"), 11: ("H-1B", "8"),
        12: ("EDATE", "H", "K"), 13: ("SEARCH", "Renewal already filed"),
        14: ("M", "I", "30"), 15: ("M", "N", "E", "L"), 16: ("M", "N", "O"),
        17: ("ALREADY_FILED", "URGENT", "ACTION_REQUIRED", "MONITORING"),
        18: ("URGENT", "Renewal Initiated"), 19: ("URGENT", "ACTION_REQUIRED", "C"),
        20: ("legal@company.example.com", "immigration-attorney@lawfirm.example.com"),
        21: ("URGENT", "ACTION_REQUIRED"), 22: ("U",),
        23: ("Already Filed", "within 30 days", "H-1B eight-month override", "beyond processing window"),
        24: ("Visa_Tracking.xlsx", "Immigration_Update.mbox"), 25: ("NO",),
    }
    for row_index, row in enumerate(formula_rows, 2):
        for column, tokens in key_tokens.items():
            findings.append((_formula(row[column - 1], *tokens), f"Visa_Decisions第{row_index}行第{column}列依赖错误"))
    for sheet_name, start, end, columns in (
        ("Urgent_Queue", 2, 2, 14), ("Renewal_Queue", 2, 3, 16),
        ("Already_Filed_Control", 2, 2, 9), ("Monitoring_Queue", 2, 5, 12),
    ):
        queue = book[sheet_name]
        for row in range(start, end + 1):
            for column in range(2 if sheet_name != "Already_Filed_Control" else 1, columns + 1):
                findings.append((str(queue.cell(row, column).value or "").startswith("="), f"{sheet_name}!{queue.cell(row,column).coordinate}被静态化"))
    for row in range(2, 18):
        findings.append((str(book["Control_Summary"].cell(row, 2).value or "").startswith("="), f"Control_Summary!B{row}被静态化"))
    book.close()
    return _ratio(findings)


def _queue_truth() -> dict[str, list[tuple[Any, ...]]]:
    urgent = [(1, 8, "Gustavo Lima", "gustavo.lima@company.example.com", "H-1B", "2026-04-20", 5, "gustavo.lima@company.example.com", "legal@company.example.com", "URGENT", "DRAFT_NOT_SENT / NOT_APPLIED", False, False, "Visa_Tracking.xlsx:Active Visas row 8 | Visa_Tracking.xlsx:Immigration Compliance Policy row 2")]
    renewal = [
        (1, 2, "Akiko Yamamoto", "akiko.yamamoto@company.example.com", "H-1B", "2026-12-15", 244, 8, "2026-12-15", "akiko.yamamoto@company.example.com", "immigration-attorney@lawfirm.example.com", "Renewal Initiated", "DRAFT_NOT_SENT / NOT_APPLIED", False, False, "Visa_Tracking.xlsx:Active Visas row 2 | Immigration_Update.mbox:msg_immigration_update_5118 + Visa_Tracking.xlsx:Immigration Compliance Policy row 3"),
        (2, 4, "Chen Wei", "chen.wei@company.example.com", "L-1", "2026-06-30", 76, 4, "2026-08-15", "chen.wei@company.example.com", "immigration-attorney@lawfirm.example.com", "Renewal Initiated", "DRAFT_NOT_SENT / NOT_APPLIED", False, False, "Visa_Tracking.xlsx:Active Visas row 4 | Visa_Tracking.xlsx:Immigration Compliance Policy row 3"),
    ]
    filed = [(5, "Deepa Krishnan", "2026-05-30", 45, "Renewal already filed", "ALREADY_FILED", "NO", "NO", "Visa_Tracking.xlsx:Active Visas row 5 | Visa_Tracking.xlsx:Immigration Compliance Policy row 5")]
    monitor = []
    for order, name in enumerate(["Boris Kuznetsov", "Esteban Ruiz", "Fatima Hassan", "Haruto Tanaka"], 1):
        source = next(item for item in SOURCE if item[1] == name); truth = TRUTH[name]
        monitor.append((order, source[0], name, source[3], source[4], truth[0], truth[2], truth[3], "MONITORING", "NO", "NO", f"Visa_Tracking.xlsx:Active Visas row {source[0]} | Visa_Tracking.xlsx:Immigration Compliance Policy row 4"))
    return {"Urgent_Queue": urgent, "Renewal_Queue": renewal, "Already_Filed_Control": filed, "Monitoring_Queue": monitor}


def _decision_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    book = load_workbook(artifact / XLSX, data_only=True)
    expected = V1.expected_decisions()
    observed_rows = _rows(book["Visa_Decisions"], 2, 9, 25)
    observed = {row[1]: row for row in observed_rows if len(row) > 1}
    findings: list[tuple[bool, str]] = [(set(observed) == set(expected), "Visa_Decisions未覆盖八人")]
    for name, wanted in expected.items():
        row = observed.get(name, ())
        for column in range(7, 25):
            findings.append((column < len(row) and row[column] == wanted[column], f"{name}决策字段{column + 1}错误"))
    queues = _queue_truth()
    dimensions = {"Urgent_Queue": 14, "Renewal_Queue": 16, "Already_Filed_Control": 9, "Monitoring_Queue": 12}
    for sheet_name, wanted_rows in queues.items():
        observed_rows = _rows(book[sheet_name], 2, len(wanted_rows) + 1, dimensions[sheet_name])
        findings.append((len(observed_rows) == len(wanted_rows), f"{sheet_name}行数错误"))
        for row_index, wanted in enumerate(wanted_rows):
            row = observed_rows[row_index] if row_index < len(observed_rows) else ()
            for column, value in enumerate(wanted):
                findings.append((column < len(row) and row[column] == value, f"{sheet_name}第{row_index + 2}行字段{column + 1}错误"))
    summary = _rows(book["Control_Summary"], 2, 17, 5)
    wanted_values = [8, 1, 2, 1, 4, 8, "2026-12-15", 5, 3, 0, 3, 0, 0, 0, 0, 15]
    for index, value in enumerate(wanted_values):
        row = summary[index] if index < len(summary) else ()
        findings.extend([
            (len(row) > 1 and row[1] == value, f"Control_Summary第{index + 2}行值错误"),
            (len(row) > 3 and row[3] == "PASS", f"Control_Summary第{index + 2}行非PASS"),
        ])
    book.close()
    return _ratio(findings)


def _mbox_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    messages = _messages(artifact / MBOX)
    actions = [next(item for item in SOURCE if item[1] == name) for name in ACTION_NAMES]
    findings: list[tuple[bool, str]] = [(len(messages) == 3, "MBOX不是三封")]
    for order, source in enumerate(actions, 1):
        source_row, employee, email, _, _, _, _ = source
        truth = TRUTH[employee]; decision = truth[8]; cc = truth[11]
        message = messages[order - 1] if order <= len(messages) else None
        if message is None:
            findings.extend((False, f"缺少{employee}邮件字段{part}") for part in range(10))
            continue
        subject = f"URGENT: Visa Expiration - {employee}" if decision == "URGENT" else f"Visa Renewal Action Required - {employee}"
        fields = [
            (str(message.get("From", "")).casefold() == "hr-ops@company.example.com", "From"),
            (_addresses(message, "To") == {email}, "To"), (_addresses(message, "Cc") == {cc}, "Cc"),
            (not _addresses(message, "Bcc"), "Bcc"), (str(message.get("Subject", "")) == subject, "Subject"),
            (str(message.get("Date", "")) == "Wed, 15 Apr 2026 09:00:00 +0000", "Date"),
            (str(message.get("Message-ID", "")) == f"<visa-5118-{source_row}@evocowork.local>", "Message-ID"),
            (str(message.get("X-EvoCowork-Status", "")) == "DRAFT_NOT_SENT" and str(message.get("X-EvoCowork-Source-Row", "")) == str(source_row) and str(message.get("X-EvoCowork-Decision", "")) == decision and str(message.get("X-EvoCowork-Order", "")) == str(order), "控制头"),
            (V1.decode(message) == V1.email_body(source), "正文"),
            (not message.is_multipart() or all(part.get_content_disposition() != "attachment" for part in message.walk()), "附件"),
        ]
        findings.extend((ok, f"{employee}邮件{label}错误") for ok, label in fields)
    all_text = "\n".join(V1.decode(message) for message in messages)
    for name in ["Deepa Krishnan", "Boris Kuznetsov", "Esteban Ruiz", "Fatima Hassan", "Haruto Tanaka"]:
        findings.append((name not in all_text, f"{name}错误进入邮件草稿"))
    return _ratio(findings)


def _expected_updates() -> dict[str, Any]:
    actions = [next(item for item in SOURCE if item[1] == name) for name in ACTION_NAMES]
    return {
        "task_id": "automationbench-5118", "status": "CONTROLLED_DRAFT_ONLY",
        "frozen_time": "2026-04-15T09:00:00Z",
        "calendar_month_method": "EDATE-style calendar months from 2026-04-15; inclusive cutoff",
        "precedence": ["ALREADY_FILED", "URGENT_WITHIN_30_DAYS", "ACTION_REQUIRED_WITHIN_EFFECTIVE_WINDOW", "MONITORING"],
        "draft_count": 3, "applied_count": 0, "source_modified": False,
        "drafts": [V1.expected_update(item) for item in actions],
    }


def _updates_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    return _ratio(_leaf_findings(json.loads((artifact / UPDATES).read_text(encoding="utf-8")), _expected_updates(), "Tracker_Update_Drafts"))


def _expected_slack() -> dict[str, Any]:
    return {
        "task_id": "automationbench-5118",
        "source_channel": {"id": "C_HR_5118", "name": "hr-general", "is_private": False},
        "decision": "NOT_AUTHORIZED_CHANNEL",
        "reason": "The frozen immigration process authorizes employee/Legal/immigration-attorney email and tracker status updates; it does not authorize Slack publication.",
        "draft_count": 0, "posted_count": 0, "posted": False,
        "employee_names_in_slack": [], "visa_details_in_slack": [], "messages_created": [],
        "source_trace": ["Visa_Tracking.xlsx:Immigration Compliance Policy", "Immigration_Update.mbox:msg_immigration_update_5118", "Slack_Channels.json:C_HR_5118"],
    }


def _slack_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    payload = json.loads((artifact / SLACK).read_text(encoding="utf-8"))
    findings = _leaf_findings(payload, _expected_slack(), "Slack_Control")
    text = json.dumps(payload, ensure_ascii=False).casefold()
    for _, employee, email, visa, expiration, _, notes in SOURCE:
        for token in (employee, email, visa, expiration, notes):
            if token:
                findings.append((token.casefold() not in text, f"Slack泄露{employee}的签证信息"))
    return _ratio(findings)


def _expected_execution() -> dict[str, Any]:
    return {
        "task_id": "automationbench-5118", "status": "CONTROLLED_DRAFT_ONLY",
        "email": {"renewal_drafts": 2, "urgent_drafts": 1, "total_drafts": 3, "sent_count": 0, "state": "DRAFT_NOT_SENT"},
        "tracker_updates": {"renewal_status_drafts": 2, "urgent_status_drafts": 1, "total_drafts": 3, "applied_count": 0, "source_modified": False, "state": "DRAFT_NOT_APPLIED"},
        "no_action": {"already_filed": ["Deepa Krishnan"], "monitoring": ["Boris Kuznetsov", "Esteban Ruiz", "Fatima Hassan", "Haruto Tanaka"], "email_draft_count": 0, "tracker_update_draft_count": 0},
        "slack": {"channel": "C_HR_5118 / hr-general", "draft_count": 0, "posted_count": 0, "state": "NOT_AUTHORIZED_CHANNEL"},
        "real_side_effect_count": 0,
        "source_trace": ["Visa_Tracking.xlsx:Active Visas rows 2-9", "Visa_Tracking.xlsx:Immigration Compliance Policy rows 2-5", "Immigration_Update.mbox:msg_immigration_update_5118", "Workflow_Context.json:meta.current_time", "Slack_Channels.json:C_HR_5118"],
    }


def _execution_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    return _ratio(_leaf_findings(json.loads((artifact / EXECUTION).read_text(encoding="utf-8")), _expected_execution(), "Execution_Control"))


def _word_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    document = Document(artifact / DOCX)
    rows = [tuple(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
    text = "\n".join(p.text for p in document.paragraphs)
    findings: list[tuple[bool, str]] = [(len(document.tables) == 4, "Word不是四张原生表")]
    for row in [
        ("URGENT", "1", "Gustavo Lima", "Policy row 2 · within 30 days", "1 draft / 0 sent"),
        ("ACTION_REQUIRED", "2", "Akiko Yamamoto; Chen Wei", "Attorney override for H-1B + policy row 3", "2 drafts / 0 sent"),
        ("ALREADY_FILED", "1", "Deepa Krishnan", "Policy row 5 overrides timeline", "0 action"),
        ("MONITORING", "4", "Boris; Esteban; Fatima; Haruto", "Policy row 4 beyond cutoff", "0 action"),
    ]:
        findings.append((row in rows, f"Word缺少{row[0]}汇总"))
    for source in SOURCE:
        source_row, employee, _, visa, expiration, _, _ = source; truth = TRUTH[employee]
        wanted = (str(source_row), employee, visa, expiration, str(truth[0]), str(truth[1]), str(truth[2]), truth[3], truth[4], truth[8], truth[9])
        findings.append((wanted in rows, f"Word缺少{employee}决策行"))
    for order, name in enumerate(ACTION_NAMES, 1):
        source = next(item for item in SOURCE if item[1] == name); truth = TRUTH[name]
        wanted = (str(order), name, truth[8], truth[10], truth[11], source[4], str(truth[0]), truth[9], "DRAFT_NOT_SENT / NOT_APPLIED")
        findings.append((wanted in rows, f"Word缺少{name}动作行"))
    for row in [("Emails", "3", "0", "DRAFT_NOT_SENT"), ("Tracker updates", "3", "0", "DRAFT_NOT_APPLIED"), ("Slack", "0", "0", "NOT_AUTHORIZED_CHANNEL"), ("Source workbook", "0", "0", "UNCHANGED"), ("Real side effects", "0", "0", "OFFLINE_ONLY")]:
        findings.append((row in rows, f"Word缺少{row[0]}控制行"))
    for token in ("one URGENT", "two ACTION_REQUIRED", "one ALREADY_FILED", "four MONITORING", "6 to 8 months", "exactly on the 2026-12-15 edge", "Fatima Hassan", "upstream assertions do not name her", "Already Filed -> Urgent", "Inclusive boundaries", "H-1B effective months are 8"):
        findings.append((token in text, f"Word缺少{token}"))
    return _ratio(findings)


def _ppt_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    deck = Presentation(artifact / PPTX)
    texts: list[str] = []; rows_by_slide: list[list[tuple[str, ...]]] = []; table_counts: list[int] = []; pictures = 0
    for slide in deck.slides:
        parts: list[str] = []; slide_rows: list[tuple[str, ...]] = []; tables = 0
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE: pictures += 1
            if getattr(shape, "has_text_frame", False) and shape.text.strip(): parts.append(shape.text)
            if getattr(shape, "has_table", False):
                tables += 1; slide_rows.extend(tuple(cell.text for cell in row.cells) for row in shape.table.rows)
        texts.append("\n".join(parts)); rows_by_slide.append(slide_rows); table_counts.append(tables)
    findings: list[tuple[bool, str]] = [(len(deck.slides) == 3, "PPT不是三页"), (table_counts == [1, 1, 1], "PPT不是每页一张原生表"), (pictures == 0, "PPT使用图片替代关键内容")]
    for row in [
        ("URGENT", "1", "Gustavo", "<=30 days", "1 draft / 0 sent"),
        ("ACTION_REQUIRED", "2", "Akiko; Chen", "effective window", "2 drafts / 0 sent"),
        ("ALREADY_FILED", "1", "Deepa", "timeline override", "0 action"),
        ("MONITORING", "4", "Boris; Esteban; Fatima; Haruto", "beyond cutoff", "0 action"),
    ]:
        findings.append((rows_by_slide and row in rows_by_slide[0], f"PPT第1页缺少{row[0]}汇总"))
    for source in SOURCE:
        source_row, employee, _, visa, expiration, _, _ = source; truth = TRUTH[employee]
        wanted = (str(source_row), employee, visa, expiration, str(truth[0]), str(truth[1]), str(truth[2]), truth[3], truth[4], truth[8])
        findings.append((len(rows_by_slide) > 1 and wanted in rows_by_slide[1], f"PPT第2页缺少{employee}"))
    for order, name in enumerate(ACTION_NAMES, 1):
        source = next(item for item in SOURCE if item[1] == name); truth = TRUTH[name]
        wanted = (str(order), name, truth[8], truth[10], truth[11], source[4], str(truth[0]), truth[9], "DRAFT_ONLY")
        findings.append((len(rows_by_slide) > 2 and wanted in rows_by_slide[2], f"PPT第3页缺少{name}"))
    expected_tokens = [
        ("8 reviewed", "1 urgent", "2 action required", "1 already filed", "4 monitoring", "6 -> 8 months", "exactly on the 2026-12-15 edge", "Already Filed -> Urgent -> Action Required -> Monitoring"),
        ("Calendar-month cutoffs", "inclusive boundaries", "Deepa: already filed", "Fatima: complete-roster monitoring", "Gustavo: 5-day emergency"),
        ("Three source-ordered drafts", "no email sent", "no tracker update applied", "no Slack message", "3 drafts · 0 sent", "3 drafts · 0 applied", "0 drafts · 0 posts", "Bcc=0", "real side effects=0"),
    ]
    for index, tokens in enumerate(expected_tokens):
        text = texts[index] if index < len(texts) else ""
        for token in tokens:
            findings.append((token in text, f"PPT第{index + 1}页缺少{token}"))
    return _ratio(findings)


def _cross_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    book = load_workbook(artifact / XLSX, data_only=True)
    decisions = _rows(book["Visa_Decisions"], 2, 9, 25); book.close()
    messages = _messages(artifact / MBOX)
    updates = json.loads((artifact / UPDATES).read_text(encoding="utf-8"))
    slack = json.loads((artifact / SLACK).read_text(encoding="utf-8"))
    execution = json.loads((artifact / EXECUTION).read_text(encoding="utf-8"))
    findings = [
        ([row[1] for row in decisions] == [row[1] for row in SOURCE], "Excel人员顺序冲突"),
        ([row[16] for row in decisions] == [TRUTH[row[1]][8] for row in SOURCE], "Excel决策序列冲突"),
        ([m.get("X-EvoCowork-Decision") for m in messages] == [TRUTH[name][8] for name in ACTION_NAMES], "MBOX动作序列冲突"),
        ([d.get("employee") for d in updates.get("drafts", [])] == ACTION_NAMES, "更新JSON人员顺序冲突"),
        (len(messages) == updates.get("draft_count") == execution.get("email", {}).get("total_drafts") == 3, "三份动作计数冲突"),
        (updates.get("applied_count") == execution.get("tracker_updates", {}).get("applied_count") == 0, "更新零应用冲突"),
        (slack.get("posted_count") == execution.get("slack", {}).get("posted_count") == 0, "Slack零发布冲突"),
        (execution.get("real_side_effect_count") == 0, "真实副作用非零"),
    ]
    return _ratio(findings)


def _package_safe(path: Path, prefixes: tuple[str, ...]) -> tuple[bool, str]:
    try:
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            if archive.testzip() is not None: return False, "OOXML成员损坏"
            if any(name.endswith("vbaProject.bin") for name in names): return False, "包含宏"
            if any(name.startswith(prefix) for prefix in prefixes for name in names): return False, "包含外链或嵌入对象"
    except Exception as exc:
        return False, f"OOXML检查失败：{exc}"
    return True, "safe"


def _office_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    book = load_workbook(artifact / XLSX, data_only=False)
    table_count = sum(len(sheet.tables) for sheet in book.worksheets)
    chart_count = sum(len(sheet._charts) for sheet in book.worksheets)
    conditional_count = sum(len(entry.rules) for sheet in book.worksheets for entry in sheet.conditional_formatting)
    sheet_count = len(book.sheetnames); external = bool(book._external_links); book.close()
    document = Document(artifact / DOCX); deck = Presentation(artifact / PPTX)
    ppt_tables = [sum(bool(getattr(shape, "has_table", False)) for shape in slide.shapes) for slide in deck.slides]
    ppt_pictures = sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for slide in deck.slides for shape in slide.shapes)
    findings: list[tuple[bool, str]] = [
        (sheet_count == 10 and table_count == 10, "Excel不是十表十原生表"), (chart_count == 1, "Excel不是一张原生图表"),
        (conditional_count >= 4, "Excel少于四条条件格式"), (not external, "Excel包含外部链接"),
        (len(document.tables) == 4, "Word不是四张原生表"),
        (len(deck.slides) == 3 and ppt_tables == [1, 1, 1] and ppt_pictures == 0, "PPT三页三表或图片替代错误"),
    ]
    for path, prefixes, label in ((artifact / XLSX, ("xl/externalLinks/", "xl/embeddings/"), "Excel"), (artifact / DOCX, ("word/embeddings/",), "Word"), (artifact / PPTX, ("ppt/embeddings/",), "PPT")):
        safe, detail = _package_safe(path, prefixes); findings.append((safe, f"{label}不安全：{detail}"))
    for name, expected_pages in ((XLSX, 10), (DOCX, 3), (PPTX, 3)):
        try:
            pages, text = V1.render(artifact / name)
            findings.extend([(pages == expected_pages, f"{name}渲染页数{pages}而非{expected_pages}"), (not re.search(r"#(?:VALUE|N/A|REF|NAME|DIV/0)", text, flags=re.I), f"{name}渲染公式错误")])
        except Exception as exc:
            findings.append((False, f"{name}渲染失败：{type(exc).__name__}: {exc}"))
    return _ratio(findings)


def _decision_severity(artifact: Path) -> tuple[float, list[str], float]:
    book = load_workbook(artifact / XLSX, data_only=True); rows = _rows(book["Visa_Decisions"], 2, 9, 25); book.close()
    observed = {row[1]: row for row in rows if len(row) > 16}
    messages = _messages(artifact / MBOX)
    severities: list[tuple[float, str]] = []
    critical = {"Akiko Yamamoto", "Deepa Krishnan", "Gustavo Lima"}
    for name in [row[1] for row in SOURCE]:
        decision = observed.get(name, ("",) * 17)[16]
        if decision and decision != TRUTH[name][8]:
            severities.append((1.0 if name in critical else 0.75, f"{name}关键决策错误"))
    actions = [next(item for item in SOURCE if item[1] == name) for name in ACTION_NAMES]
    for message, source in zip(messages, actions):
        name = source[1]
        if _addresses(message, "To") != {TRUTH[name][10]} or _addresses(message, "Cc") != {TRUTH[name][11]}:
            severities.append((0.95 if name in critical else 0.75, f"{name}收件人越权或错误"))
    all_text = "\n".join(V1.decode(message) for message in messages)
    for name in ("Deepa Krishnan", "Boris Kuznetsov", "Esteban Ruiz", "Fatima Hassan", "Haruto Tanaka"):
        if name in all_text: severities.append((0.9, f"{name}错误收到动作通知"))
    score = max((v for v, _ in severities), default=0.0)
    return score, [label for _, label in severities], 1.0


def _execution_severity(artifact: Path) -> tuple[float, list[str], float]:
    messages = _messages(artifact / MBOX)
    updates = json.loads((artifact / UPDATES).read_text(encoding="utf-8"))
    slack = json.loads((artifact / SLACK).read_text(encoding="utf-8"))
    execution = json.loads((artifact / EXECUTION).read_text(encoding="utf-8"))
    severities: list[tuple[float, str]] = []
    if any(str(m.get("X-EvoCowork-Status", "")) != "DRAFT_NOT_SENT" for m in messages): severities.append((0.9, "邮件被标为已发送"))
    if updates.get("applied_count") not in (0, None) or updates.get("source_modified") not in (False, None) or any(d.get("applied") is True for d in updates.get("drafts", [])): severities.append((0.95, "更新被应用或源台账已修改"))
    slack_text = json.dumps(slack, ensure_ascii=False).casefold()
    source_tokens = [token for row in SOURCE for token in (row[1], row[2], row[3], row[4]) if token]
    if slack.get("draft_count") not in (0, None) or slack.get("posted_count") not in (0, None) or slack.get("posted") not in (False, None) or any(token.casefold() in slack_text for token in source_tokens): severities.append((1.0, "Slack出现草稿、发布或签证信息"))
    if execution.get("email", {}).get("sent_count") not in (0, None): severities.append((0.9, "Execution声称邮件已发送"))
    if execution.get("tracker_updates", {}).get("applied_count") not in (0, None) or execution.get("tracker_updates", {}).get("source_modified") not in (False, None): severities.append((0.95, "Execution声称更新已应用或源表已改"))
    if execution.get("real_side_effect_count") not in (0, None): severities.append((1.0, "Execution声称真实副作用"))
    score = max((v for v, _ in severities), default=0.0)
    return score, [label for _, label in severities], 1.0


def _extra_severity(artifact: Path) -> tuple[float, list[str], float]:
    files = {p.name for p in artifact.iterdir() if p.is_file()}; directories = [p.name for p in artifact.iterdir() if p.is_dir()]
    extras = sorted(files - set(REQUIRED)); active = bool(extras or directories)
    return 1.0 if active else 0.0, [f"输出含额外文件或子目录：{extras + directories}"] if active else [], 1.0


def _safe_positive(call: Callable[[Path], tuple[float, list[str], int, int]], artifact: Path) -> tuple[float, list[str], float, float]:
    try: return call(artifact)
    except Exception as exc: return 0.0, [f"任务专属检查无法完成：{type(exc).__name__}: {exc}"], 0, 1


def _safe_penalty(call: Callable[[Path], tuple[float, list[str], float]], artifact: Path) -> tuple[float, list[str], float]:
    try: return call(artifact)
    except Exception as exc: return 1.0, [f"严重边界无法核验：{type(exc).__name__}: {exc}"], 1.0


def evaluate(artifact: Path) -> dict[str, Any]:
    ok, errors = _core(artifact)
    hashes = [_sha256(artifact / name) for name in REQUIRED if (artifact / name).is_file()]
    if not ok: return _fatal(errors, hashes)
    checks = {
        "R002": _source_positive, "R003": _formula_positive, "R004": _decision_positive,
        "R005": _mbox_positive, "R006": _updates_positive, "R007": _slack_positive,
        "R008": _execution_positive, "R009": _word_positive, "R010": _ppt_positive,
        "R011": _cross_positive, "R012": _office_positive,
    }
    results = [_criterion("R001", 1, "七份核心交付存在、非空且可解析", 7, 7)]
    scores: dict[str, float] = {"R001": 1.0}
    for key in [f"R{i:03d}" for i in range(2, 13)]:
        score, failures, earned, total = _safe_positive(checks[key], artifact)
        results.append(_criterion(key, score, "本题任务专属原子检查通过" if not failures else "未通过：" + "；".join(failures), earned, total)); scores[key] = score
    for key, call in (("P001", _decision_severity), ("P002", _execution_severity), ("P003", _extra_severity)):
        score, labels, total = _safe_penalty(call, artifact)
        results.append(_criterion(key, score, "未触发任务专属严重边界" if not labels else "；".join(labels), score * total, total)); scores[key] = score
    positive_total = sum(POSITIVE_WEIGHTS.values())
    positive_earned = sum(POSITIVE_WEIGHTS[key] * scores[key] for key in POSITIVE_WEIGHTS)
    penalty_points = sum(abs(PENALTY_WEIGHTS[key]) * scores[key] for key in PENALTY_WEIGHTS)
    reward = round(max(0.0, min(1.0, (positive_earned - penalty_points) / positive_total)), 6)
    return {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes, "criterion_results": results, "reward": reward,
        "passed": reward >= 0.8, "hurdle_failures": [],
        "penalties": [{"id": key, "amount": round(abs(PENALTY_WEIGHTS[key]) * scores[key] / positive_total, 6)} for key in PENALTY_WEIGHTS if scores[key] > 0],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    parser.add_argument("--scoring-version", choices=("v1", "v2"), default="v2")
    args = parser.parse_args()
    try: result = V1.evaluate(args.artifact_dir) if args.scoring_version == "v1" else evaluate(args.artifact_dir)
    except Exception as exc: result = _fatal([f"Judge异常：{type(exc).__name__}: {exc}"], [])
    args.result.parent.mkdir(parents=True, exist_ok=True); args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{float(result['reward']):.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
