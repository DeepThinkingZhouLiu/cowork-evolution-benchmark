#!/usr/bin/env python3
"""只把 AutomationBench 5119 的公开初始状态静态化为本题输入。

本脚本只接受 ``hr.expense_report_audit``，逐项断言八条待审费用、完整
费用政策邮件、冻结时间、空 Finance Slack 快照和十条上游 assertions。
它不是通用制题器，不会创造员工、金额、收据、例外、规则或动作。
"""

from __future__ import annotations

import hashlib
import importlib.util
import json
import shutil
import sys
import types
import zipfile
from datetime import datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
UPSTREAM = Path(__file__).with_name("upstream-hr-tasks.py")
SNAPSHOT = Path(__file__).with_name("upstream-task-5119.json")
LICENSE = Path(__file__).with_name("LICENSE-AutomationBench-MIT")
DATA_DIR = ROOT / "data" / "input_files"
ENV_DIR = ROOT / "environment" / "input_files"
FIXED_ZIP_TIME = (2026, 4, 15, 9, 0, 0)
SOURCE_SHA256 = "c173a36a03f43096bd122577584aa32f9df379d8760aa46c43be04c85d101779"
LICENSE_SHA256 = "c6f9d87aff6d627d5cd78cb09364dafed7d6f1c4bb7a91183d57377157e96ba4"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_task() -> dict:
    if sha256(UPSTREAM) != SOURCE_SHA256 or sha256(LICENSE) != LICENSE_SHA256:
        raise ValueError("冻结的 AutomationBench 源文件或 LICENSE 哈希错误")
    datasets = types.ModuleType("datasets")
    datasets.Dataset = object
    sys.modules["datasets"] = datasets
    for name in ("automationbench", "automationbench.domains", "automationbench.domains.hr"):
        sys.modules[name] = types.ModuleType(name)
    noise = types.ModuleType("automationbench.domains.hr._noise")
    noise.apply_noise = lambda value: value
    sys.modules[noise.__name__] = noise
    spec = importlib.util.spec_from_file_location("automationbench_hr_tasks_frozen_5119", UPSTREAM)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入冻结的 AutomationBench HR 源文件")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    task = module.get_task_5119()
    if task.get("example_id") != 5119 or task.get("task") != "hr.expense_report_audit":
        raise ValueError("冻结函数不再是 AutomationBench 5119")
    return task


def assert_shape(task: dict) -> None:
    prompt = task["prompt"][1]["content"]
    assert all(
        token in prompt
        for token in (
            "approve all the pending expense reports",
            "Update the tracker and send confirmation emails",
            "under $100",
            "don't bother with receipt verification",
            "follow our expense policy",
        )
    )
    state = task["info"]["initial_state"]
    assert state["meta"] == {"schema_version": "0.1.0", "current_time": "2026-04-15T09:00:00Z"}
    messages = state["gmail"]["messages"]
    assert [item["id"] for item in messages] == ["msg_expense_policy_5119"]
    policy = messages[0]
    assert policy["from_"] == "cfo@company.example.com"
    assert policy["subject"] == "2026 Expense Policy"
    assert all(
        token in policy["body_plain"]
        for token in (
            "Meals: up to $75",
            "Business class requires VP approval",
            "Software/Tools: up to $500",
            "All expenses over $1,000 require receipt",
            "Missing receipts: leave the report unchanged and do not notify the employee",
            "every expense — regardless of amount — must have a receipt",
            "strict no-exceptions receipt policy",
        )
    )
    spreadsheet = state["google_sheets"]["spreadsheets"][0]
    assert spreadsheet["id"] == "ss_expenses_5119"
    assert spreadsheet["title"] == "Expense Report Tracker"
    worksheet = spreadsheet["worksheets"][0]
    assert worksheet["id"] == "ws_pending_5119" and worksheet["title"] == "Pending Reports"
    rows = worksheet["rows"]
    assert [item["row_id"] for item in rows] == list(range(2, 10))
    assert [item["cells"]["Employee"] for item in rows] == [
        "Alice Johansson",
        "Ben Torres",
        "Claudia Reeves",
        "David Osei",
        "Elena Marchetti",
        "Finn Larsson",
        "Grace Pham",
        "Hector Kim",
    ]
    assert [item["cells"]["Amount"] for item in rows] == [
        "$62.00",
        "$120.00",
        "$2,400.00",
        "$350.00",
        "$1,800.00",
        "$799.00",
        "$95.00",
        "$450.00",
    ]
    assert [item["cells"]["Receipt"] for item in rows] == ["Yes", "Yes", "Yes", "Yes", "No", "Yes", "Yes", "Yes"]
    assert rows[2]["cells"]["Description"] == "Flight to NYC - business class"
    assert rows[4]["cells"]["Description"] == "Conference flights + hotel"
    assert rows[6]["cells"]["Notes"] == "Pre-approved by CFO on April 3"
    assert state["slack"] == {"channels": [{"id": "C_FIN_5119", "name": "finance", "is_private": False}], "messages": []}
    assertions = task["info"]["assertions"]
    assert len(assertions) == 10
    serialized = json.dumps(assertions)
    assert all(name in serialized for name in ("Ben Torres", "Claudia Reeves", "Finn Larsson"))
    assert "elena.marchetti@company.example.com" in serialized
    assert not any(name in serialized for name in ("Grace Pham", "Hector Kim"))


def normalize_xlsx(path: Path) -> None:
    temporary = path.with_suffix(".normalized.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    temporary.replace(path)


def style_sheet(sheet, widths: dict[str, float]) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = sheet.dimensions


def write_tracker(state: dict, target: Path) -> None:
    source = state["google_sheets"]["spreadsheets"][0]
    worksheet = source["worksheets"][0]
    workbook = Workbook()
    workbook.creator = "AutomationBench 5119 public fixture"
    workbook.title = source["title"]
    workbook.subject = "Static Office representation of the pending expense report tracker"
    workbook.description = "No employee, amount, receipt, note or expense fact was added beyond AutomationBench 5119."
    sheet = workbook.active
    sheet.title = worksheet["title"]
    headers = ["Source Row", "Employee", "Employee Email", "Category", "Amount", "Description", "Receipt", "Status", "Notes"]
    sheet.append(headers)
    for item in worksheet["rows"]:
        cells = item["cells"]
        sheet.append(
            [
                item["row_id"],
                cells["Employee"],
                cells["Employee Email"],
                cells["Category"],
                cells["Amount"],
                cells["Description"],
                cells["Receipt"],
                cells["Status"],
                cells["Notes"],
            ]
        )
    table = Table(displayName="PendingExpenseSourceInput", ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    style_sheet(sheet, {"A": 12, "B": 25, "C": 39, "D": 19, "E": 14, "F": 38, "G": 12, "H": 15, "I": 34})
    for row in range(2, 10):
        sheet.row_dimensions[row].height = 30
    workbook.save(target)
    normalize_xlsx(target)


def message_bytes(record: dict) -> bytes:
    message = EmailMessage(policy=SMTP)
    message["From"] = record["from_"]
    message["To"] = ", ".join(record["to"])
    message["Subject"] = record["subject"]
    parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
    message["Date"] = parsed.strftime("%a, %d %b %Y %H:%M:%S +0000")
    message["Message-ID"] = f"<{record['id']}@automationbench.invalid>"
    message["X-AutomationBench-Message-ID"] = record["id"]
    message["X-AutomationBench-Thread-ID"] = record["thread_id"]
    message["X-AutomationBench-Labels"] = ",".join(record.get("label_ids", []))
    message.set_content(record["body_plain"], charset="utf-8")
    buffer = BytesIO()
    BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
    return buffer.getvalue()


def write_mbox(messages: list[dict], target: Path) -> None:
    chunks: list[bytes] = []
    for record in messages:
        parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
        chunks.append(f"From {record['from_']} {parsed.strftime('%a %b %d %H:%M:%S %Y')}\n".encode("ascii"))
        chunks.append(message_bytes(record).replace(b"\r\n", b"\n"))
        chunks.append(b"\n")
    target.write_bytes(b"".join(chunks))


def main() -> None:
    task = load_task()
    assert_shape(task)
    state = task["info"]["initial_state"]
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    ENV_DIR.mkdir(parents=True, exist_ok=True)
    SNAPSHOT.write_text(json.dumps(task, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_tracker(state, DATA_DIR / "Expense_Report_Tracker.xlsx")
    write_mbox(state["gmail"]["messages"], DATA_DIR / "Expense_Policy.mbox")
    (DATA_DIR / "Slack_Channels.json").write_text(json.dumps(state["slack"], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (DATA_DIR / "Workflow_Context.json").write_text(json.dumps({"meta": state["meta"]}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    for source in sorted(DATA_DIR.iterdir()):
        if source.is_file():
            shutil.copy2(source, ENV_DIR / source.name)


if __name__ == "__main__":
    main()
