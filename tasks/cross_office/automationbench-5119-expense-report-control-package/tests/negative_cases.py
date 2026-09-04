#!/usr/bin/env python3
"""AutomationBench 5119 的任务专属 v2 分层负例。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from typing import Callable

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
INPUTS = ROOT / "environment" / "input_files"
OUTPUTS = (
    "Expense_Report_Audit.xlsx",
    "Expense_Control_Brief.docx",
    "Expense_Review.pptx",
    "Expense_Notification_Drafts.mbox",
    "Expense_Tracker_Update_Drafts.json",
    "Expense_Slack_Control.json",
    "Expense_Execution_Control.json",
)
Mutator = Callable[[Path], None]


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"无法载入{path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _xlsx(root: Path, sheet: str, cell: str, value) -> None:
    path = root / OUTPUTS[0]
    book = load_workbook(path)
    book[sheet][cell] = value
    book.save(path)


def _json(root: Path, name: str, mutator: Callable[[dict], None]) -> None:
    path = root / name
    payload = json.loads(path.read_text(encoding="utf-8"))
    mutator(payload)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _rewrite_mbox(root: Path, mutator: Callable[[list], None]) -> None:
    path = root / OUTPUTS[3]
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    mutator(messages)
    with path.open("wb") as handle:
        for index, message in enumerate(messages, 1):
            handle.write(f"From MAILER-DAEMON Wed Apr 15 09:00:{index:02d} 2026\n".encode())
            BytesGenerator(handle, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
            handle.write(b"\n")


def _replace_word(root: Path, old: str, new: str) -> None:
    path = root / OUTPUTS[1]
    document = Document(path)
    changed = False
    for paragraph in document.paragraphs:
        if old in paragraph.text:
            paragraph.text = paragraph.text.replace(old, new, 1)
            changed = True
            break
    if not changed:
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if old in cell.text:
                        cell.text = cell.text.replace(old, new, 1)
                        changed = True
                        break
                if changed:
                    break
            if changed:
                break
    if not changed:
        raise ValueError(f"Word中找不到{old!r}")
    document.save(path)


def _replace_ppt(root: Path, old: str, new: str) -> None:
    path = root / OUTPUTS[2]
    deck = Presentation(path)
    changed = False
    for slide in deck.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and old in shape.text:
                shape.text = shape.text.replace(old, new, 1)
                changed = True
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        if old in cell.text:
                            cell.text = cell.text.replace(old, new, 1)
                            changed = True
                            break
                    if changed:
                        break
            if changed:
                break
        if changed:
            break
    if not changed:
        raise ValueError(f"PPT中找不到{old!r}")
    deck.save(path)


def _input_only(root: Path) -> None:
    for path in list(root.iterdir()):
        if path.is_file():
            path.unlink()
    for source in INPUTS.iterdir():
        if source.is_file():
            shutil.copy2(source, root / source.name)


def _cases() -> list[tuple[str, Mutator, str]]:
    cases: list[tuple[str, Mutator, str]] = []
    for name in OUTPUTS:
        cases.append((f"missing-{name}", lambda root, item=name: (root / item).unlink(), "fatal_zero"))
        cases.append((f"corrupt-{name}", lambda root, item=name: (root / item).write_bytes(b"not-a-valid-artifact"), "fatal_zero"))
    cases.extend([
        ("mbox-empty", lambda root: (root / OUTPUTS[3]).write_bytes(b""), "fatal_zero"),
        ("unmodified-input-only", _input_only, "fatal_zero"),
        ("source-alice-amount-local", lambda root: _xlsx(root, "Source_Reports", "E2", "$63.00"), "partial"),
        ("alice-amount-formula-static", lambda root: _xlsx(root, "Expense_Decisions", "F2", 62), "partial"),
        ("elena-missing-receipt-bypass", lambda root: _xlsx(root, "Expense_Decisions", "O6", "APPROVED"), "partial"),
        ("grace-preapproval-lost", lambda root: _xlsx(root, "Expense_Decisions", "L8", "NO"), "partial"),
        ("ben-route-role-wrong", lambda root: _xlsx(root, "Expense_Decisions", "N3", "VP Travel"), "partial"),
        ("mail-alice-description-local", lambda root: _rewrite_mbox(root, lambda messages: messages[0].set_payload(messages[0].get_payload().replace("Client dinner", "Client meal"))), "partial"),
        ("mail-claudia-wrong-recipient", lambda root: _rewrite_mbox(root, lambda messages: messages[2].replace_header("To", "cfo@company.example.com")), "partial"),
        ("update-finn-amount-local", lambda root: _json(root, OUTPUTS[4], lambda payload: payload["drafts"][4].update({"amount": "$798.00"})), "partial"),
        ("slack-expense-leak", lambda root: _json(root, OUTPUTS[5], lambda payload: payload.update({"draft_count": 1, "messages_created": [{"text": "Elena Marchetti $1,800.00 receipt missing"}]})), "partial"),
        ("execution-email-sent", lambda root: _json(root, OUTPUTS[6], lambda payload: payload["email"].update({"sent_count": 1, "state": "SENT"})), "partial"),
        ("updates-applied", lambda root: _json(root, OUTPUTS[4], lambda payload: payload.update({"applied_count": 1, "source_modified": True})), "partial"),
        ("word-elena-policy-local", lambda root: _replace_word(root, "Elena Marchetti remains Pending", "Elena Marchetti may be approved"), "partial"),
        ("ppt-grace-omitted", lambda root: _replace_ppt(root, "Grace Pham", "Unlisted employee"), "partial"),
        ("ppt-policy-local", lambda root: _replace_ppt(root, "every expense requires a receipt", "some expenses require a receipt"), "partial"),
        ("control-approved-count-conflict", lambda root: _xlsx(root, "Control_Summary", "B5", 3), "partial"),
        ("extra-file", lambda root: (root / "unexpected.txt").write_text("unexpected", encoding="utf-8"), "partial"),
    ])
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    judge = _load("expense_5119_v2_judge", ROOT / "tests" / "judge.py")
    judge.V1.render = lambda path: ({".xlsx": 10, ".docx": 3, ".pptx": 3}[path.suffix.casefold()], "render-ok")
    if judge.evaluate(GOLDEN).get("reward") != 1.0:
        raise AssertionError("Golden v2必须为1.0")
    results: list[dict] = []
    for name, mutator, layer in _cases():
        with tempfile.TemporaryDirectory(prefix=f"expense-5119-v2-{name}-") as temporary:
            output = Path(temporary) / "output"
            shutil.copytree(GOLDEN, output)
            mutator(output)
            result = judge.evaluate(output)
            reward = float(result.get("reward", 0))
            if layer == "fatal_zero" and reward != 0:
                raise AssertionError(f"{name}应硬归零，实际{reward}")
            if layer == "partial" and not (0 < reward < 1):
                raise AssertionError(f"{name}应保留连续部分分，实际{reward}")
            results.append({
                "case": name,
                "layer": layer,
                "reward": reward,
                "passed": bool(result.get("passed")),
                "penalties": result.get("penalties", []),
            })
    partial = [item["reward"] for item in results if item["layer"] == "partial"]
    report = {
        "total": len(results),
        "fatal_zero": sum(item["layer"] == "fatal_zero" for item in results),
        "partial_continuous": len(partial),
        "partial_reward_range": [min(partial), max(partial)],
        "cases": results,
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
