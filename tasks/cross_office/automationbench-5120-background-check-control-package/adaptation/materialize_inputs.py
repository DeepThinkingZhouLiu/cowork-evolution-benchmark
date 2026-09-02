#!/usr/bin/env python3
"""只把 AutomationBench 5120 的公开初始状态静态化为本题输入。

本脚本只接受 ``hr.background_check_followup``，逐项断言八条候选人台账、
供应商更正、Legal 处理政策、冻结时间、Recruiting Slack 快照和九条上游
assertions。它不是通用制题器，不会创造候选人、结果、政策或动作。
"""

from __future__ import annotations

import hashlib
import importlib.util
import json
import shutil
import sys
import types
import zipfile
from datetime import datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
UPSTREAM = Path(__file__).with_name("upstream-hr-tasks.py")
SNAPSHOT = Path(__file__).with_name("upstream-task-5120.json")
LICENSE = Path(__file__).with_name("LICENSE-AutomationBench-MIT")
DATA_DIR = ROOT / "data" / "input_files"
ENV_DIR = ROOT / "environment" / "input_files"
FIXED_ZIP_TIME = (2026, 4, 15, 9, 0, 0)
SOURCE_SHA256 = "c173a36a03f43096bd122577584aa32f9df379d8760aa46c43be04c85d101779"
LICENSE_SHA256 = "c6f9d87aff6d627d5cd78cb09364dafed7d6f1c4bb7a91183d57377157e96ba4"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_task() -> dict:
    if sha256(UPSTREAM) != SOURCE_SHA256 or sha256(LICENSE) != LICENSE_SHA256:
        raise ValueError("冻结的 AutomationBench 源文件或 LICENSE 哈希错误")
    datasets = types.ModuleType("datasets")
    datasets.Dataset = object
    sys.modules["datasets"] = datasets
    for name in ("automationbench", "automationbench.domains", "automationbench.domains.hr"):
        sys.modules[name] = types.ModuleType(name)
    noise = types.ModuleType("automationbench.domains.hr._noise")
    noise.apply_noise = lambda value: value
    sys.modules[noise.__name__] = noise
    spec = importlib.util.spec_from_file_location("automationbench_hr_tasks_frozen_5120", UPSTREAM)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入冻结的 AutomationBench HR 源文件")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    task = module.get_task_5120()
    if task.get("example_id") != 5120 or task.get("task") != "hr.background_check_followup":
        raise ValueError("冻结函数不再是 AutomationBench 5120")
    return task


def assert_shape(task: dict) -> None:
    prompt = task["prompt"][1]["content"]
    assert all(token in prompt for token in ("pending background check results", "appropriate next step", "background check procedures", "Google Sheets and Gmail", "#recruiting", "names of affected entities"))
    state = task["info"]["initial_state"]
    assert state["meta"] == {"schema_version": "0.1.0", "current_time": "2026-04-15T09:00:00Z"}
    messages = state["gmail"]["messages"]
    assert [item["id"] for item in messages] == ["msg_bgcheck_correction_5120", "msg_bgcheck_policy_5120"]
    correction, policy = messages
    assert correction["from_"] == "screening@bgcheck-vendor.example.com"
    assert correction["subject"] == "Updated Background Check Result - Naomi Yoshida"
    assert all(token in correction["body_plain"] for token in ("Naomi Yoshida", "BGC-4412", "changed from 'Clear' to 'Conditional'", "pending court disposition", "does not disqualify employment", "legal team notification"))
    assert policy["from_"] == "legal@company.example.com"
    assert policy["subject"] == "Background Check Result Handling Procedures"
    assert all(token in policy["body_plain"] for token in ("Clear: Proceed with onboarding", "Failed: Rescind offer", "Do NOT email the candidate", "Conditional: Proceed with onboarding BUT notify Legal", "Pending: No action"))

    spreadsheet = state["google_sheets"]["spreadsheets"][0]
    assert spreadsheet["id"] == "ss_bgcheck_5120" and spreadsheet["title"] == "Background Check Tracker"
    worksheet = spreadsheet["worksheets"][0]
    assert worksheet["id"] == "ws_checks_5120" and worksheet["title"] == "Candidates"
    rows = worksheet["rows"]
    assert [item["row_id"] for item in rows] == list(range(2, 10))
    assert [item["cells"]["Candidate"] for item in rows] == ["Liam Chen", "Maya Okafor", "Naomi Yoshida", "Oliver Strand", "Petra Johansson", "Quinn Davis", "Rosa Martinez", "Sanjay Gupta"]
    assert [item["cells"]["Candidate ID"] for item in rows] == ["BGC-4401", "BGC-4405", "BGC-4412", "BGC-4418", "BGC-4422", "BGC-4430", "BGC-4435", "BGC-4441"]
    assert [item["cells"]["Result"] for item in rows] == ["Clear", "Failed", "Clear", "Pending", "Clear", "Failed", "Clear", "Pending"]
    assert [item["cells"]["Status"] for item in rows] == ["Awaiting Action"] * 8
    assert state["slack"] == {"channels": [{"id": "C_REC_5120", "name": "recruiting", "is_private": False}], "messages": []}
    assertions = task["info"]["assertions"]
    assert len(assertions) == 9
    serialized = json.dumps(assertions)
    assert all(token in serialized for token in ("Liam Chen", "Maya Okafor", "Naomi Yoshida", "Quinn Davis", "Conditional - Legal Notified", "recruiting"))


def normalize_xlsx(path: Path) -> None:
    temporary = path.with_suffix(".normalized.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    temporary.replace(path)


def style_sheet(sheet, widths: dict[str, float]) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = sheet.dimensions


def write_tracker(state: dict, target: Path) -> None:
    source = state["google_sheets"]["spreadsheets"][0]
    worksheet = source["worksheets"][0]
    workbook = Workbook()
    workbook.creator = "AutomationBench 5120 public fixture"
    workbook.title = source["title"]
    workbook.subject = "Static Office representation of the background check tracker"
    workbook.description = "No candidate, result, manager, position or status was added beyond AutomationBench 5120."
    sheet = workbook.active
    sheet.title = worksheet["title"]
    headers = ["Source Row", "Candidate", "Candidate ID", "Hiring Manager", "Manager Email", "Position", "Result", "Status", "Notes"]
    sheet.append(headers)
    for item in worksheet["rows"]:
        cells = item["cells"]
        sheet.append([item["row_id"], cells["Candidate"], cells["Candidate ID"], cells["Hiring Manager"], cells["Manager Email"], cells["Position"], cells["Result"], cells["Status"], cells["Notes"]])
    table = Table(displayName="BackgroundCheckSourceInput", ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    style_sheet(sheet, {"A":12,"B":24,"C":16,"D":22,"E":38,"F":27,"G":15,"H":22,"I":28})
    for row in range(2, 10):
        sheet.row_dimensions[row].height = 27
    workbook.save(target)
    normalize_xlsx(target)


def message_bytes(record: dict) -> bytes:
    message = EmailMessage(policy=SMTP)
    message["From"] = record["from_"]
    message["To"] = ", ".join(record["to"])
    message["Subject"] = record["subject"]
    parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
    message["Date"] = parsed.strftime("%a, %d %b %Y %H:%M:%S +0000")
    message["Message-ID"] = f"<{record['id']}@automationbench.invalid>"
    message["X-AutomationBench-Message-ID"] = record["id"]
    message["X-AutomationBench-Thread-ID"] = record["thread_id"]
    message["X-AutomationBench-Labels"] = ",".join(record.get("label_ids", []))
    message.set_content(record["body_plain"], charset="utf-8")
    buffer = BytesIO()
    BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
    return buffer.getvalue()


def write_mbox(messages: list[dict], target: Path) -> None:
    chunks: list[bytes] = []
    for record in messages:
        parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
        chunks.append(f"From {record['from_']} {parsed.strftime('%a %b %d %H:%M:%S %Y')}\n".encode("ascii"))
        chunks.append(message_bytes(record).replace(b"\r\n", b"\n"))
        chunks.append(b"\n")
    target.write_bytes(b"".join(chunks))


def main() -> None:
    task = load_task()
    assert_shape(task)
    state = task["info"]["initial_state"]
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    ENV_DIR.mkdir(parents=True, exist_ok=True)
    SNAPSHOT.write_text(json.dumps(task, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_tracker(state, DATA_DIR / "Background_Check_Tracker.xlsx")
    write_mbox(state["gmail"]["messages"], DATA_DIR / "Background_Check_Messages.mbox")
    (DATA_DIR / "Slack_Channels.json").write_text(json.dumps(state["slack"], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (DATA_DIR / "Workflow_Context.json").write_text(json.dumps({"meta": state["meta"]}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    for source in sorted(DATA_DIR.iterdir()):
        if source.is_file():
            shutil.copy2(source, ENV_DIR / source.name)


if __name__ == "__main__":
    main()
