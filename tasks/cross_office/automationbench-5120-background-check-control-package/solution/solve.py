#!/usr/bin/env python3
"""AutomationBench 5120 背景调查跟进任务的逐题参考解法。

这不是通用制题器：输入、八名候选人、两封证据邮件、覆盖优先级、四类
处理规则、通知路由和隐私边界均逐项冻结为本题真值。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import mailbox
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt as PptPt


INPUT_HASHES = {
    "Background_Check_Messages.mbox": "d7645da7106062a38838f28b957e3c0d6cabc0e4fb560f9b432e2620736a2c68",
    "Background_Check_Tracker.xlsx": "b29b0e70e736c6dfd945741d006028be6cd46e6e5840a00fcbf4adf73c0a60ba",
    "Slack_Channels.json": "568ea5e7aeaa300bf8e509d4c725ad43d6e799069dc671e615df7038b62cd951",
    "Workflow_Context.json": "76c4a0657bc27f2fd830b10613ec7ceae3adcdc5798aeee0f822d27bf69bbc99",
}
XLSX_NAME = "Background_Check_Audit.xlsx"
DOCX_NAME = "Background_Check_Control_Brief.docx"
PPTX_NAME = "Background_Check_Review.pptx"
MBOX_NAME = "Background_Check_Notification_Drafts.mbox"
UPDATES_NAME = "Background_Check_Tracker_Update_Drafts.json"
SLACK_NAME = "Background_Check_Slack_Drafts.json"
EXECUTION_NAME = "Background_Check_Execution_Control.json"
EXPECTED_OUTPUTS = {XLSX_NAME, DOCX_NAME, PPTX_NAME, MBOX_NAME, UPDATES_NAME, SLACK_NAME, EXECUTION_NAME}
FIXED_ISO_TIME = "2026-04-15T09:00:00Z"
FIXED_ZIP_TIME = (2026, 4, 15, 9, 0, 0)

NAVY = "17365D"
BLUE = "5B9BD5"
GREEN = "E2F0D9"
AMBER = "FFF2CC"
RED = "F4CCCC"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="B7C9D6")
PPT_NAVY = PptRGBColor(23, 54, 93)
PPT_BLUE = PptRGBColor(91, 155, 213)
PPT_LIGHT_BLUE = PptRGBColor(221, 235, 247)
PPT_GREEN = PptRGBColor(226, 240, 217)
PPT_AMBER = PptRGBColor(255, 242, 204)
PPT_RED = PptRGBColor(244, 204, 204)
PPT_WHITE = PptRGBColor(255, 255, 255)


@dataclass(frozen=True)
class CandidateRecord:
    source_row: int
    candidate: str
    candidate_id: str
    hiring_manager: str
    manager_email: str
    position: str
    source_result: str
    source_status: str
    notes: str


@dataclass(frozen=True)
class Decision:
    source: CandidateRecord
    vendor_override: str
    effective_result: str
    decision: str
    proposed_status: str
    notify_to: str
    email_draft: str
    tracker_update: str
    slack_label: str
    privacy_control: str
    authority: str


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def decode_message(message) -> str:
    if message.is_multipart():
        parts = []
        for part in message.walk():
            if part.get_content_type() == "text/plain":
                payload = part.get_payload(decode=True) or b""
                parts.append(payload.decode(part.get_content_charset() or "utf-8", errors="replace"))
        return "\n".join(parts).strip()
    payload = message.get_payload(decode=True) or b""
    return payload.decode(message.get_content_charset() or "utf-8", errors="replace").strip()


def read_inputs(input_dir: Path) -> tuple[list[CandidateRecord], dict[str, dict[str, str]], dict, str]:
    for name, expected in INPUT_HASHES.items():
        source = input_dir / name
        if not source.is_file() or sha256(source) != expected:
            raise ValueError(f"输入 {name} 缺失或 SHA256 与冻结版本不一致")

    context = json.loads((input_dir / "Workflow_Context.json").read_text(encoding="utf-8"))
    if context != {"meta": {"schema_version": "0.1.0", "current_time": FIXED_ISO_TIME}}:
        raise ValueError("冻结运行时间已变化")
    channels = json.loads((input_dir / "Slack_Channels.json").read_text(encoding="utf-8"))
    expected_channels = {"channels": [{"id": "C_REC_5120", "name": "recruiting", "is_private": False}], "messages": []}
    if channels != expected_channels:
        raise ValueError("Recruiting Slack 快照已变化")

    workbook = load_workbook(input_dir / "Background_Check_Tracker.xlsx", read_only=True, data_only=True)
    if workbook.sheetnames != ["Candidates"]:
        raise ValueError("背景调查台账结构已变化")
    records: list[CandidateRecord] = []
    for row in workbook["Candidates"].iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        records.append(CandidateRecord(int(row[0]), *(str(value or "") for value in row[1:9])))
    workbook.close()
    expected_records = [
        (2, "Liam Chen", "BGC-4401", "Derek Huang", "derek.huang@company.example.com", "Software Engineer", "Clear", "Awaiting Action", ""),
        (3, "Maya Okafor", "BGC-4405", "Mira Patel", "mira.patel@company.example.com", "Marketing Analyst", "Failed", "Awaiting Action", ""),
        (4, "Naomi Yoshida", "BGC-4412", "Lisa Wang", "lisa.wang@company.example.com", "Account Executive", "Clear", "Awaiting Action", ""),
        (5, "Oliver Strand", "BGC-4418", "Derek Huang", "derek.huang@company.example.com", "DevOps Engineer", "Pending", "Awaiting Action", ""),
        (6, "Petra Johansson", "BGC-4422", "Mira Patel", "mira.patel@company.example.com", "Content Writer", "Clear", "Awaiting Action", ""),
        (7, "Quinn Davis", "BGC-4430", "Lisa Wang", "lisa.wang@company.example.com", "Sales Rep", "Failed", "Awaiting Action", ""),
        (8, "Rosa Martinez", "BGC-4435", "Derek Huang", "derek.huang@company.example.com", "QA Engineer", "Clear", "Awaiting Action", ""),
        (9, "Sanjay Gupta", "BGC-4441", "Mira Patel", "mira.patel@company.example.com", "Data Analyst", "Pending", "Awaiting Action", ""),
    ]
    actual_records = [
        (item.source_row, item.candidate, item.candidate_id, item.hiring_manager, item.manager_email, item.position, item.source_result, item.source_status, item.notes)
        for item in records
    ]
    if actual_records != expected_records:
        raise ValueError("八条冻结候选人记录已变化")

    box = mailbox.mbox(input_dir / "Background_Check_Messages.mbox", create=False)
    raw_messages = list(box)
    box.close()
    if len(raw_messages) != 2:
        raise ValueError("背景调查证据 MBOX 必须恰好两封")
    messages: dict[str, dict[str, str]] = {}
    for message in raw_messages:
        message_id = str(message.get("X-AutomationBench-Message-ID", ""))
        messages[message_id] = {
            "message_id": message_id,
            "thread_id": str(message.get("X-AutomationBench-Thread-ID", "")),
            "from": str(message.get("From", "")),
            "to": str(message.get("To", "")),
            "subject": str(message.get("Subject", "")),
            "date": str(message.get("Date", "")),
            "body": decode_message(message),
        }
    if list(messages) != ["msg_bgcheck_correction_5120", "msg_bgcheck_policy_5120"]:
        raise ValueError("两封证据邮件顺序或身份已变化")
    correction = messages["msg_bgcheck_correction_5120"]
    policy = messages["msg_bgcheck_policy_5120"]
    if correction["from"].casefold() != "screening@bgcheck-vendor.example.com" or correction["subject"] != "Updated Background Check Result - Naomi Yoshida":
        raise ValueError("供应商更正邮件身份已变化")
    if not all(token in correction["body"] for token in ("Naomi Yoshida", "BGC-4412", "changed from 'Clear' to 'Conditional'", "pending court disposition", "does not disqualify employment", "legal team notification")):
        raise ValueError("供应商更正证据不完整")
    if policy["from"].casefold() != "legal@company.example.com" or policy["subject"] != "Background Check Result Handling Procedures":
        raise ValueError("Legal 政策邮件身份已变化")
    if not all(token in policy["body"] for token in ("Clear: Proceed with onboarding", "Failed: Rescind offer", "Do NOT email the candidate", "Conditional: Proceed with onboarding BUT notify Legal", "Pending: No action")):
        raise ValueError("Legal 政策证据不完整")
    return records, messages, channels, context["meta"]["current_time"]


def derive_decisions(records: list[CandidateRecord]) -> list[Decision]:
    decisions: list[Decision] = []
    for source in records:
        override = "Conditional" if source.candidate_id == "BGC-4412" else ""
        effective = override or source.source_result
        if effective == "Clear":
            values = (
                "CLEARED_PROCEED", "Cleared - Proceed", source.manager_email, "YES", "YES",
                "ONBOARDING_MAY_PROCEED", "Manager receives clearance confirmation; Slack receives operational label only.",
                "Legal procedure 1: Clear; email the hiring manager and propose Cleared - Proceed.",
            )
        elif effective == "Failed":
            values = (
                "FAILED_RESCIND", "Failed - Offer Rescinded", "hr-director@company.example.com", "YES", "YES",
                "HR_DIRECTOR_ACTION", "HR Director only; never email candidate; Slack omits raw result.",
                "Legal procedure 2: rescind offer, notify HR Director, and never email the candidate directly.",
            )
        elif effective == "Conditional":
            values = (
                "CONDITIONAL_LEGAL", "Conditional - Legal Notified", "legal@company.example.com", "YES", "YES",
                "LEGAL_NOTIFICATION_BEFORE_START", "Legal only; no manager notice; sensitive vendor reason omitted outside restricted evidence.",
                "Vendor correction msg_bgcheck_correction_5120 overrides tracker row 4; Legal procedure 3 applies.",
            )
        elif effective == "Pending":
            values = (
                "PENDING_NO_ACTION", source.source_status, "", "NO", "NO",
                "PENDING_NO_ACTION", "No email or tracker update; Slack receives operational label only.",
                "Legal procedure 4: Pending requires no action and source status remains unchanged.",
            )
        else:
            raise ValueError(f"未冻结的背景调查结果：{effective}")
        decisions.append(Decision(source, override, effective, *values))

    expected = {
        "Liam Chen": ("", "Clear", "CLEARED_PROCEED", "Cleared - Proceed", "derek.huang@company.example.com", "YES", "YES", "ONBOARDING_MAY_PROCEED"),
        "Maya Okafor": ("", "Failed", "FAILED_RESCIND", "Failed - Offer Rescinded", "hr-director@company.example.com", "YES", "YES", "HR_DIRECTOR_ACTION"),
        "Naomi Yoshida": ("Conditional", "Conditional", "CONDITIONAL_LEGAL", "Conditional - Legal Notified", "legal@company.example.com", "YES", "YES", "LEGAL_NOTIFICATION_BEFORE_START"),
        "Oliver Strand": ("", "Pending", "PENDING_NO_ACTION", "Awaiting Action", "", "NO", "NO", "PENDING_NO_ACTION"),
        "Petra Johansson": ("", "Clear", "CLEARED_PROCEED", "Cleared - Proceed", "mira.patel@company.example.com", "YES", "YES", "ONBOARDING_MAY_PROCEED"),
        "Quinn Davis": ("", "Failed", "FAILED_RESCIND", "Failed - Offer Rescinded", "hr-director@company.example.com", "YES", "YES", "HR_DIRECTOR_ACTION"),
        "Rosa Martinez": ("", "Clear", "CLEARED_PROCEED", "Cleared - Proceed", "derek.huang@company.example.com", "YES", "YES", "ONBOARDING_MAY_PROCEED"),
        "Sanjay Gupta": ("", "Pending", "PENDING_NO_ACTION", "Awaiting Action", "", "NO", "NO", "PENDING_NO_ACTION"),
    }
    actual = {
        item.source.candidate: (item.vendor_override, item.effective_result, item.decision, item.proposed_status, item.notify_to, item.email_draft, item.tracker_update, item.slack_label)
        for item in decisions
    }
    if actual != expected:
        raise ValueError("八人背景调查决策真值不一致")
    return decisions


POLICY_CLAUSES = [
    (1, "Clear", "Proceed with onboarding; email hiring manager; propose Cleared - Proceed."),
    (2, "Failed", "Rescind offer; email HR Director; propose Failed - Offer Rescinded; do not email candidate."),
    (3, "Conditional", "Proceed with onboarding but notify Legal before start; propose Conditional - Legal Notified."),
    (4, "Pending", "No action; leave Status unchanged."),
]


def source_trace(row: int) -> str:
    suffix = " | Background_Check_Messages.mbox:msg_bgcheck_correction_5120" if row == 4 else ""
    return f"Background_Check_Tracker.xlsx:Candidates row {row}{suffix} | Background_Check_Messages.mbox:msg_bgcheck_policy_5120"


def add_table(sheet, name: str) -> None:
    table = Table(displayName=name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)


def style_sheet(sheet, widths: dict[str, float], footer: str, *, a3: bool = False) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", bold=True, color=WHITE, size=8)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)
    sheet.row_dimensions[1].height = 31
    grid = Side(style="hair", color="D9E2F3")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=7.7)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=grid, right=grid, bottom=grid)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if a3 else sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.oddFooter.center.text = footer
    sheet.print_area = sheet.dimensions


def normalize_ooxml(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="background-5120-ooxml-") as temporary:
        target = Path(temporary) / path.name
        axis_ids: dict[str, str] = {}
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as destination:
            for name in sorted(source.namelist()):
                data = source.read(name)
                if name == "docProps/core.xml":
                    try:
                        root = etree.fromstring(data)
                        ns = {"dcterms": "http://purl.org/dc/terms/", "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"}
                        for tag in ("created", "modified"):
                            for node in root.xpath(f"//dcterms:{tag}", namespaces=ns):
                                node.text = FIXED_ISO_TIME
                        for tag in ("lastModifiedBy", "revision"):
                            for node in root.xpath(f"//cp:{tag}", namespaces=ns):
                                node.text = "EvoCowork" if tag == "lastModifiedBy" else "1"
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except etree.XMLSyntaxError:
                        pass
                if name.startswith("xl/charts/chart") and name.endswith(".xml"):
                    try:
                        root = etree.fromstring(data)
                        ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
                        for node in root.xpath("//c:axId", namespaces=ns):
                            old = node.get("val")
                            if old and old not in axis_ids:
                                axis_ids[old] = str(51_200_001 + len(axis_ids))
                        for node in root.xpath("//c:axId | //c:crossAx", namespaces=ns):
                            old = node.get("val")
                            if old in axis_ids:
                                node.set("val", axis_ids[old])
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except etree.XMLSyntaxError:
                        pass
                info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                destination.writestr(info, data)
        shutil.copy2(target, path)


def recalculate_xlsx(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="background-5120-recalc-") as temporary:
        root = Path(temporary)
        source_dir, output_dir, profile_dir = root / "source", root / "output", root / "profile"
        source_dir.mkdir()
        output_dir.mkdir()
        source = source_dir / path.name
        shutil.copy2(path, source)
        process = subprocess.run(
            ["libreoffice", "--headless", f"-env:UserInstallation=file://{profile_dir}", "--convert-to", "xlsx", "--outdir", str(output_dir), str(source)],
            stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300,
        )
        converted = output_dir / path.name
        if process.returncode != 0 or not converted.is_file():
            raise RuntimeError(f"LibreOffice 公式重算失败：{process.stdout[-1000:]}")
        shutil.copy2(converted, path)


def write_workbook(path: Path, records: list[CandidateRecord], decisions: list[Decision], messages: dict[str, dict[str, str]], channels: dict, frozen_time: str) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.creator = "EvoCowork AutomationBench 5120"
    workbook.title = "Background Check Audit"
    workbook.subject = "Vendor override, Legal routing and privacy-controlled offline actions"
    workbook.description = "Eight source candidates, exact message evidence, formula-derived decisions and draft-only controls."
    fixed_dt = datetime(2026, 4, 15, 9, 0, 0)
    workbook.created = fixed_dt
    workbook.modified = fixed_dt

    source = workbook.create_sheet("Source_Candidates")
    source.append(["Source Row", "Candidate", "Candidate ID", "Hiring Manager", "Manager Email", "Position", "Result", "Status", "Notes"])
    for item in records:
        source.append([item.source_row, item.candidate, item.candidate_id, item.hiring_manager, item.manager_email, item.position, item.source_result, item.source_status, item.notes])
    add_table(source, "BackgroundSourceCandidates")
    style_sheet(source, {"A":11,"B":24,"C":16,"D":22,"E":38,"F":27,"G":14,"H":21,"I":24}, "RESTRICTED · AutomationBench 5120 source", a3=True)

    message_sheet = workbook.create_sheet("Message_Source")
    message_sheet.append(["Message ID", "From", "To", "Subject", "Date", "Full Body", "Authority", "Sensitive Handling"])
    correction = messages["msg_bgcheck_correction_5120"]
    policy = messages["msg_bgcheck_policy_5120"]
    message_sheet.append([correction["message_id"], correction["from"], correction["to"], correction["subject"], correction["date"], correction["body"], "Overrides Naomi tracker result", "Restricted evidence; detailed reason excluded from all outbound drafts and open-channel summary"])
    message_sheet.append([policy["message_id"], policy["from"], policy["to"], policy["subject"], policy["date"], policy["body"], "Governs four result classes", "Operational policy; candidate-result details still minimized by audience"])
    add_table(message_sheet, "BackgroundMessageSource")
    style_sheet(message_sheet, {"A":34,"B":40,"C":37,"D":48,"E":30,"F":120,"G":42,"H":72}, "Exact vendor correction and Legal procedure", a3=True)
    message_sheet.row_dimensions[2].height = 145
    message_sheet.row_dimensions[3].height = 205

    evidence = workbook.create_sheet("Evidence_Register")
    evidence.append(["Evidence ID", "Type", "Subject", "Exact Fact", "Use", "Privacy", "Source Trace"])
    for item in records:
        fact = f"candidate={item.candidate}; id={item.candidate_id}; manager={item.hiring_manager}; manager_email={item.manager_email}; position={item.position}; result={item.source_result}; status={item.source_status}; notes={item.notes}"
        evidence.append([f"EV-ROW-{item.source_row}", "CANDIDATE_ROW", item.candidate, fact, "Base record", "Restricted HR data", f"Background_Check_Tracker.xlsx:Candidates row {item.source_row}"])
    evidence.append(["EV-OVERRIDE-1", "VENDOR_CORRECTION", "Naomi Yoshida / BGC-4412", correction["body"], "Override Clear with Conditional", "Detailed reason restricted", "Background_Check_Messages.mbox:msg_bgcheck_correction_5120"])
    for number, result, clause in POLICY_CLAUSES:
        evidence.append([f"EV-POLICY-{number}", "LEGAL_POLICY", result, clause, "Decision mapping", "Audience-specific disclosure", "Background_Check_Messages.mbox:msg_bgcheck_policy_5120"])
    evidence.append(["EV-TIME-1", "TIME", "Frozen runtime", frozen_time, "Artifact timestamp", "Non-sensitive", "Workflow_Context.json:meta.current_time"])
    channel = channels["channels"][0]
    evidence.append(["EV-SLACK-1", "CHANNEL", f"{channel['id']} / {channel['name']}", "is_private=false; messages=[]", "One offline minimal-disclosure draft", "Open channel: names + safe operational labels only", "Slack_Channels.json:C_REC_5120"])
    evidence.append(["EV-REQUEST-1", "USER_REQUEST", "Background-check follow-up", "Use tracker and Gmail, take appropriate next step, update #recruiting with affected names", "Defines requested workflow", "Policy and privacy still govern content", "AutomationBench 5120 prompt"])
    add_table(evidence, "BackgroundEvidenceRegister")
    style_sheet(evidence, {"A":18,"B":23,"C":34,"D":112,"E":42,"F":48,"G":68}, "16 source-backed evidence records", a3=True)

    decision_sheet = workbook.create_sheet("Candidate_Decisions")
    headers = [
        "Source Row", "Candidate", "Candidate ID", "Hiring Manager", "Manager Email", "Position", "Source Result", "Source Status", "Notes",
        "Vendor Override", "Effective Result", "Is Clear", "Is Failed", "Is Conditional", "Is Pending", "Decision", "Proposed Status", "Notify To",
        "Email Draft", "Tracker Update", "Slack Label", "Privacy Control", "Decision Basis", "Source Trace",
    ]
    decision_sheet.append(headers)
    for index in range(2, 10):
        decision_sheet.append([
            f"=Source_Candidates!A{index}", f"=Source_Candidates!B{index}", f"=Source_Candidates!C{index}", f"=Source_Candidates!D{index}", f"=Source_Candidates!E{index}",
            f"=Source_Candidates!F{index}", f"=Source_Candidates!G{index}", f"=Source_Candidates!H{index}", f"=Source_Candidates!I{index}",
            f'=IF(C{index}="BGC-4412","Conditional","")', f'=IF(J{index}<>"",J{index},G{index})',
            f'=IF(K{index}="Clear","YES","NO")', f'=IF(K{index}="Failed","YES","NO")', f'=IF(K{index}="Conditional","YES","NO")', f'=IF(K{index}="Pending","YES","NO")',
            f'=IF(L{index}="YES","CLEARED_PROCEED",IF(M{index}="YES","FAILED_RESCIND",IF(N{index}="YES","CONDITIONAL_LEGAL",IF(O{index}="YES","PENDING_NO_ACTION","ERROR"))))',
            f'=IF(L{index}="YES","Cleared - Proceed",IF(M{index}="YES","Failed - Offer Rescinded",IF(N{index}="YES","Conditional - Legal Notified",IF(O{index}="YES",H{index},"ERROR"))))',
            f'=IF(L{index}="YES",E{index},IF(M{index}="YES","hr-director@company.example.com",IF(N{index}="YES","legal@company.example.com","")))',
            f'=IF(O{index}="YES","NO","YES")', f'=IF(O{index}="YES","NO","YES")',
            f'=IF(L{index}="YES","ONBOARDING_MAY_PROCEED",IF(M{index}="YES","HR_DIRECTOR_ACTION",IF(N{index}="YES","LEGAL_NOTIFICATION_BEFORE_START","PENDING_NO_ACTION")))',
            f'=IF(L{index}="YES","Manager receives clearance confirmation; Slack receives operational label only.",IF(M{index}="YES","HR Director only; never email candidate; Slack omits raw result.",IF(N{index}="YES","Legal only; no manager notice; sensitive vendor reason omitted outside restricted evidence.","No email or tracker update; Slack receives operational label only.")))',
            f'=IF(L{index}="YES","Legal procedure 1: Clear; email the hiring manager and propose Cleared - Proceed.",IF(M{index}="YES","Legal procedure 2: rescind offer, notify HR Director, and never email the candidate directly.",IF(N{index}="YES","Vendor correction msg_bgcheck_correction_5120 overrides tracker row 4; Legal procedure 3 applies.","Legal procedure 4: Pending requires no action and source status remains unchanged.")))',
            f'="Background_Check_Tracker.xlsx:Candidates row "&A{index}&IF(A{index}=4," | Background_Check_Messages.mbox:msg_bgcheck_correction_5120","")&" | Background_Check_Messages.mbox:msg_bgcheck_policy_5120"',
        ])
    add_table(decision_sheet, "BackgroundDecisionMatrix")
    widths = {chr(65 + i): width for i, width in enumerate([10,22,16,21,37,26,14,20,18,17,16,11,11,15,12,21,31,39,13,15,34,70,82,90])}
    style_sheet(decision_sheet, widths, "Formula-derived background-check decisions · vendor correction has precedence", a3=True)
    for column, value, color in (("P", "CLEARED_PROCEED", GREEN), ("P", "FAILED_RESCIND", RED), ("P", "CONDITIONAL_LEGAL", AMBER), ("P", "PENDING_NO_ACTION", GRAY)):
        decision_sheet.conditional_formatting.add(f"{column}2:{column}9", FormulaRule(formula=[f'{column}2="{value}"'], fill=PatternFill("solid", fgColor=color)))
    decision_sheet.conditional_formatting.add("J2:J9", FormulaRule(formula=['J2<>""'], fill=PatternFill("solid", fgColor=AMBER)))

    clear = workbook.create_sheet("Clear_Queue")
    clear.append(["Order", "Source Row", "Candidate", "Candidate ID", "Hiring Manager", "Notify To", "Decision", "Proposed Status", "State", "Source Trace"])
    for order, item in enumerate((entry for entry in decisions if entry.effective_result == "Clear"), 1):
        source_item = item.source
        clear.append([order, source_item.source_row, source_item.candidate, source_item.candidate_id, source_item.hiring_manager, item.notify_to, item.decision, item.proposed_status, "DRAFT_ONLY", source_trace(source_item.source_row)])
    add_table(clear, "BackgroundClearQueue")
    style_sheet(clear, {"A":9,"B":11,"C":24,"D":16,"E":22,"F":39,"G":22,"H":28,"I":18,"J":85}, "3 clearance routes · manager draft only", a3=True)

    failed = workbook.create_sheet("Failed_Control")
    failed.append(["Order", "Source Row", "Candidate", "Candidate ID", "Hiring Manager", "Manager Notified", "Notify To", "Candidate Direct Email", "Decision", "Proposed Status", "Slack Label", "State", "Source Trace"])
    for order, item in enumerate((entry for entry in decisions if entry.effective_result == "Failed"), 1):
        source_item = item.source
        failed.append([order, source_item.source_row, source_item.candidate, source_item.candidate_id, source_item.hiring_manager, "NO", item.notify_to, "PROHIBITED", item.decision, item.proposed_status, item.slack_label, "DRAFT_ONLY", source_trace(source_item.source_row)])
    add_table(failed, "BackgroundFailedControl")
    style_sheet(failed, {"A":9,"B":11,"C":24,"D":16,"E":22,"F":18,"G":40,"H":24,"I":22,"J":31,"K":28,"L":18,"M":84}, "2 failed results · HR Director only · candidate direct email prohibited", a3=True)

    conditional = workbook.create_sheet("Conditional_Control")
    conditional.append(["Source Row", "Candidate", "Candidate ID", "Tracker Result", "Vendor Override", "Effective Result", "Hiring Manager", "Manager Notified", "Notify To", "Decision", "Proposed Status", "Sensitive Reason Outbound", "Slack Label", "State", "Source Trace"])
    item = next(entry for entry in decisions if entry.effective_result == "Conditional")
    source_item = item.source
    conditional.append([source_item.source_row, source_item.candidate, source_item.candidate_id, source_item.source_result, item.vendor_override, item.effective_result, source_item.hiring_manager, "NO", item.notify_to, item.decision, item.proposed_status, "OMITTED", item.slack_label, "DRAFT_ONLY", source_trace(source_item.source_row)])
    add_table(conditional, "BackgroundConditionalControl")
    style_sheet(conditional, {"A":11,"B":24,"C":16,"D":16,"E":18,"F":18,"G":22,"H":18,"I":38,"J":23,"K":31,"L":27,"M":38,"N":18,"O":96}, "Naomi vendor override · Legal only · sensitive reason omitted", a3=True)

    pending = workbook.create_sheet("Pending_Control")
    pending.append(["Source Row", "Candidate", "Candidate ID", "Source Result", "Source Status", "Effective Result", "Decision", "Proposed Status", "Email Draft", "Tracker Update", "Slack Label", "State", "Source Trace"])
    for item in (entry for entry in decisions if entry.effective_result == "Pending"):
        source_item = item.source
        pending.append([source_item.source_row, source_item.candidate, source_item.candidate_id, source_item.source_result, source_item.source_status, item.effective_result, item.decision, item.proposed_status, item.email_draft, item.tracker_update, item.slack_label, "NO_ACTION", source_trace(source_item.source_row)])
    add_table(pending, "BackgroundPendingControl")
    style_sheet(pending, {"A":11,"B":24,"C":16,"D":16,"E":22,"F":18,"G":23,"H":22,"I":13,"J":15,"K":29,"L":17,"M":84}, "Oliver and Sanjay · status unchanged · no email or tracker draft", a3=True)

    communication = workbook.create_sheet("Communication_Register")
    communication.append(["Order", "Source Row", "Candidate", "Effective Result", "Decision", "Email To", "Email State", "Tracker State", "Slack Operational Label", "Slack State", "Privacy Control", "Source Trace"])
    for order, item in enumerate(decisions, 1):
        source_item = item.source
        communication.append([order, source_item.source_row, source_item.candidate, item.effective_result, item.decision, item.notify_to, "DRAFT_NOT_SENT" if item.email_draft == "YES" else "NO_DRAFT", "DRAFT_NOT_APPLIED" if item.tracker_update == "YES" else "NO_DRAFT", item.slack_label, "DRAFT_NOT_POSTED", item.privacy_control, source_trace(source_item.source_row)])
    add_table(communication, "BackgroundCommunicationRegister")
    style_sheet(communication, {"A":9,"B":11,"C":24,"D":18,"E":23,"F":39,"G":22,"H":24,"I":38,"J":23,"K":78,"L":90}, "6 email drafts · 6 tracker drafts · 1 Slack draft · zero execution", a3=True)

    summary = workbook.create_sheet("Control_Summary")
    summary.append(["Metric", "Actual", "Expected", "State", "Source Trace"])
    summary_rows = [
        ("Source candidates", 8, 8, "COMPLETE", "Source_Candidates rows 2-9"),
        ("Effective Clear", 3, 3, "COMPLETE", "Candidate_Decisions:Is Clear"),
        ("Effective Failed", 2, 2, "COMPLETE", "Candidate_Decisions:Is Failed"),
        ("Effective Conditional", 1, 1, "OVERRIDE_APPLIED", "Candidate_Decisions:Vendor Override"),
        ("Effective Pending", 2, 2, "NO_ACTION", "Candidate_Decisions:Is Pending"),
        ("Email drafts", 6, 6, "DRAFT_NOT_SENT", MBOX_NAME),
        ("Emails sent", 0, 0, "NONE", EXECUTION_NAME),
        ("Tracker update drafts", 6, 6, "DRAFT_NOT_APPLIED", UPDATES_NAME),
        ("Tracker updates applied", 0, 0, "NONE", EXECUTION_NAME),
        ("Slack drafts", 1, 1, "DRAFT_NOT_POSTED", SLACK_NAME),
        ("Slack posts", 0, 0, "NONE", EXECUTION_NAME),
        ("Source modifications", 0, 0, "UNCHANGED", EXECUTION_NAME),
        ("Real side effects", 0, 0, "OFFLINE_ONLY", EXECUTION_NAME),
    ]
    for row in summary_rows:
        summary.append(row)
    add_table(summary, "BackgroundControlSummary")
    style_sheet(summary, {"A":30,"B":14,"C":14,"D":28,"E":70}, "Cross-artifact control totals · all external actions remain drafts")
    chart = BarChart()
    chart.type = "col"
    chart.title = "Effective result mix"
    chart.y_axis.title = "Candidates"
    chart.height = 6.5
    chart.width = 11.5
    chart.add_data(Reference(summary, min_col=2, min_row=3, max_row=6), titles_from_data=False)
    chart.set_categories(Reference(summary, min_col=1, min_row=3, max_row=6))
    chart.legend = None
    summary.add_chart(chart, "G2")
    summary.print_area = "A1:R14"

    workbook.save(path)
    normalize_ooxml(path)
    recalculate_xlsx(path)
    normalize_ooxml(path)


def set_cell_shading(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def add_doc_table(document: Document, headers: list[str], rows: list[tuple], widths: list[float], font_size: float = 7.0) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = value
        cell.width = Cm(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(font_size)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    for row_number, values in enumerate(rows, 1):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
            cells[index].width = Cm(widths[index])
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_number % 2:
                set_cell_shading(cells[index], "DDEBF7")
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(font_size)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_doc_title(document: Document, title: str, subtitle: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(title)
    run.font.name = "Aptos Display"
    run.font.size = Pt(21)
    run.font.bold = True
    run.font.color.rgb = RGBColor(23, 54, 93)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(subtitle)
    run.font.name = "Aptos"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(89, 89, 89)


def add_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.03
    for run in paragraph.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8.5)


def write_brief(path: Path, decisions: list[Decision]) -> None:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.1)
    section.bottom_margin = Cm(1.1)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)
    document.core_properties.author = "EvoCowork AutomationBench 5120"
    document.core_properties.title = "Background Check Control Brief"
    document.core_properties.created = datetime(2026, 4, 15, 9, 0, 0)
    document.core_properties.modified = datetime(2026, 4, 15, 9, 0, 0)

    add_doc_title(document, "Background Check Control Brief", "AutomationBench 5120 · vendor correction precedence · restricted draft only")
    add_body(document, "Eight source candidates resolve to three CLEARED_PROCEED, two FAILED_RESCIND, one CONDITIONAL_LEGAL and two PENDING_NO_ACTION. The vendor's 2026-04-14 correction overrides Naomi Yoshida's tracker result from Clear to Conditional. Her detailed vendor reason remains restricted: Legal is notified before start, Lisa Wang is not notified, and the open recruiting-channel draft uses only a safe operational label.")
    add_doc_table(document, ["Effective class", "Count", "Candidates", "Route", "Proposed status", "Execution"], [
        ("Clear", 3, "Liam Chen; Petra Johansson; Rosa Martinez", "respective hiring manager", "Cleared - Proceed", "3 drafts / 0 sent"),
        ("Failed", 2, "Maya Okafor; Quinn Davis", "HR Director only", "Failed - Offer Rescinded", "2 drafts / 0 sent; candidate direct email prohibited"),
        ("Conditional", 1, "Naomi Yoshida", "Legal only; no manager", "Conditional - Legal Notified", "1 draft / 0 sent; reason omitted"),
        ("Pending", 2, "Oliver Strand; Sanjay Gupta", "no email / no update", "Awaiting Action", "0 action"),
    ], [3.5,1.7,7.2,5.4,5.8,7.2], font_size=6.6)
    add_body(document, "Legal mapping is deterministic: Clear -> manager; Failed -> HR Director and never candidate; Conditional -> Legal before start; Pending -> no action. All six email and six tracker actions remain offline drafts. One #recruiting draft contains all eight names in source order and only minimized operational labels.")

    document.add_page_break()
    add_doc_title(document, "Complete eight-candidate decision register", "Every source row retains the original result, effective result, route, status and privacy boundary")
    rows = []
    for item in decisions:
        source = item.source
        rows.append((source.source_row, source.candidate, source.candidate_id, source.hiring_manager, source.position, source.source_result, item.vendor_override or "—", item.effective_result, item.decision, item.proposed_status, item.notify_to or "—"))
    add_doc_table(document, ["Row", "Candidate", "ID", "Manager", "Position", "Source", "Override", "Effective", "Decision", "Proposed status", "Notify to"], rows, [1.1,3.1,2.2,3.0,3.3,1.7,1.9,2.0,3.2,4.2,5.3], font_size=5.7)
    add_body(document, "Naomi's source row still says Clear, but the effective result is Conditional because msg_bgcheck_correction_5120 is later, candidate-specific evidence. Oliver Strand and Sanjay Gupta remain Awaiting Action and generate neither email nor tracker-update drafts. Maya Okafor and Quinn Davis route only to hr-director@company.example.com; no candidate-direct message exists.")

    document.add_page_break()
    add_doc_title(document, "Controlled communications and privacy boundary", "Six emails · six tracker updates · one Slack draft · zero real execution")
    actions = [item for item in decisions if item.email_draft == "YES"]
    add_doc_table(document, ["Order", "Candidate", "Effective", "Decision", "To", "Proposed status", "State"], [
        (order, item.source.candidate, item.effective_result, item.decision, item.notify_to, item.proposed_status, "DRAFT_NOT_SENT / NOT_APPLIED")
        for order, item in enumerate(actions, 1)
    ], [1.2,3.6,2.5,3.5,6.5,5.0,6.0], font_size=6.2)
    add_doc_table(document, ["Control", "Draft / proposed", "Executed", "State"], [
        ("Emails", 6, 0, "DRAFT_NOT_SENT"), ("Tracker updates", 6, 0, "DRAFT_NOT_APPLIED"),
        ("Recruiting Slack", 1, 0, "DRAFT_NOT_POSTED"), ("Source workbook", 0, 0, "UNCHANGED"), ("Real side effects", 0, 0, "OFFLINE_ONLY"),
    ], [7.0,5.0,4.0,10.0], font_size=7.0)
    add_body(document, "The recruiting draft lists Liam Chen, Maya Okafor, Naomi Yoshida, Oliver Strand, Petra Johansson, Quinn Davis, Rosa Martinez and Sanjay Gupta in source order. Its labels are only ONBOARDING_MAY_PROCEED, HR_DIRECTOR_ACTION, LEGAL_NOTIFICATION_BEFORE_START and PENDING_NO_ACTION. Candidate IDs, emails, positions, raw Failed/Conditional result words and the vendor's detailed reason are excluded.")
    document.save(path)
    normalize_ooxml(path)


def add_textbox(slide, x, y, w, h, text, *, font_size=16, bold=False, color=PPT_NAVY, align=PP_ALIGN.LEFT, fill=None, line=None, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill or PPT_WHITE
    shape.line.color.rgb = line or (fill or PPT_WHITE)
    frame = shape.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = Inches(0.07)
    frame.margin_right = Inches(0.07)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.03)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = PptPt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_slide_title(slide, title, subtitle):
    add_textbox(slide, 0.55, 0.25, 12.25, 0.55, title, font_size=24, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER)
    add_textbox(slide, 0.75, 0.78, 11.85, 0.36, subtitle, font_size=10.5, color=PptRGBColor(89, 89, 89), align=PP_ALIGN.CENTER)


def style_ppt_table(table, body_font=8.5):
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.fill.solid()
            cell.fill.fore_color.rgb = PPT_NAVY if row_index == 0 else (PPT_LIGHT_BLUE if row_index % 2 else PPT_WHITE)
            cell.margin_left = Inches(0.04)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.LEFT
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = PptPt(9 if row_index == 0 else body_font)
                    run.font.bold = row_index == 0
                    run.font.color.rgb = PPT_WHITE if row_index == 0 else PPT_NAVY


def write_presentation(path: Path, decisions: list[Decision]) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]
    presentation.core_properties.author = "EvoCowork AutomationBench 5120"
    presentation.core_properties.title = "Background Check Review"
    presentation.core_properties.created = datetime(2026, 4, 15, 9, 0, 0)
    presentation.core_properties.modified = datetime(2026, 4, 15, 9, 0, 0)

    slide = presentation.slides.add_slide(blank)
    add_slide_title(slide, "Background check control review", "8 candidates · vendor correction precedence · all actions remain offline drafts")
    cards = [("3", "clear / manager route", PPT_GREEN), ("2", "failed / HR Director", PPT_RED), ("1", "conditional / Legal", PPT_AMBER), ("2", "pending / no action", PPT_LIGHT_BLUE)]
    for index, (value, label, fill) in enumerate(cards):
        left = 0.7 + index * 3.15
        add_textbox(slide, left, 1.25, 2.75, 0.72, value, font_size=25, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER, fill=fill, line=PPT_BLUE, radius=True)
        add_textbox(slide, left, 1.92, 2.75, 0.34, label, font_size=9, color=PPT_NAVY, align=PP_ALIGN.CENTER)
    table = slide.shapes.add_table(5, 6, Inches(0.45), Inches(2.48), Inches(12.45), Inches(2.75)).table
    rows = [
        ("Class", "Count", "Candidates", "Route", "Proposed status", "Execution"),
        ("Clear", 3, "Liam; Petra; Rosa", "respective manager", "Cleared - Proceed", "3 drafts / 0 sent"),
        ("Failed", 2, "Maya; Quinn", "HR Director only", "Failed - Offer Rescinded", "2 drafts / 0 sent"),
        ("Conditional", 1, "Naomi", "Legal only", "Conditional - Legal Notified", "1 draft / 0 sent"),
        ("Pending", 2, "Oliver; Sanjay", "no email / update", "Awaiting Action", "0 action"),
    ]
    for row_index, values in enumerate(rows):
        for column, value in enumerate(values):
            table.cell(row_index, column).text = str(value)
    for index, width in enumerate([1.35,0.7,2.55,2.15,3.25,2.45]):
        table.columns[index].width = Inches(width)
    style_ppt_table(table, 7.8)
    add_textbox(slide, 0.55, 5.5, 12.2, 0.65, "Naomi: tracker Clear -> vendor correction Conditional -> Legal only before start", font_size=13, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER, fill=PPT_AMBER, line=PPT_AMBER, radius=True)
    add_textbox(slide, 0.55, 6.32, 12.2, 0.45, "Sensitive vendor reason stays restricted · Lisa Wang is not notified · open-channel draft uses safe labels only", font_size=10.2, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER, fill=PPT_RED, line=PPT_RED, radius=True)
    add_textbox(slide, 0.55, 6.92, 12.2, 0.28, "Email sent=0 · tracker applied=0 · Slack posted=0 · real side effects=0", font_size=9.4, color=PPT_NAVY, align=PP_ALIGN.CENTER)

    slide = presentation.slides.add_slide(blank)
    add_slide_title(slide, "Complete eight-candidate decision register", "Source result · vendor override · effective result · controlled route")
    table = slide.shapes.add_table(9, 10, Inches(0.2), Inches(1.18), Inches(12.95), Inches(5.85)).table
    headers = ("Row", "Candidate", "ID", "Manager", "Position", "Source", "Override", "Effective", "Decision", "Notify to")
    for column, value in enumerate(headers):
        table.cell(0, column).text = value
    for row_index, item in enumerate(decisions, 1):
        source = item.source
        values = (source.source_row, source.candidate, source.candidate_id, source.hiring_manager, source.position, source.source_result, item.vendor_override or "—", item.effective_result, item.decision, item.notify_to or "—")
        for column, value in enumerate(values):
            table.cell(row_index, column).text = str(value)
    for index, width in enumerate([0.5,1.45,1.05,1.25,1.5,0.75,0.95,0.85,1.45,3.2]):
        table.columns[index].width = Inches(width)
    style_ppt_table(table, 6.7)
    add_textbox(slide, 0.42, 6.94, 12.48, 0.27, "No source row omitted · Pending rows unchanged · Failed candidates never contacted directly", font_size=9.2, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER, fill=PPT_LIGHT_BLUE, line=PPT_LIGHT_BLUE, radius=True)

    slide = presentation.slides.add_slide(blank)
    add_slide_title(slide, "Controlled communications", "Six source-ordered email and tracker drafts · one minimized recruiting draft · zero execution")
    actions = [item for item in decisions if item.email_draft == "YES"]
    table = slide.shapes.add_table(7, 7, Inches(0.25), Inches(1.18), Inches(12.85), Inches(3.72)).table
    headers = ("Order", "Candidate", "Effective", "Decision", "To", "Proposed status", "State")
    for column, value in enumerate(headers):
        table.cell(0, column).text = value
    for row_index, item in enumerate(actions, 1):
        values = (row_index, item.source.candidate, item.effective_result, item.decision, item.notify_to, item.proposed_status, "DRAFT_ONLY")
        for column, value in enumerate(values):
            table.cell(row_index, column).text = str(value)
    for index, width in enumerate([0.6,1.55,1.0,1.65,3.0,3.15,1.9]):
        table.columns[index].width = Inches(width)
    style_ppt_table(table, 6.8)
    add_textbox(slide, 0.55, 5.18, 3.8, 0.72, "Email\n6 drafts · 0 sent", font_size=13, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER, fill=PPT_GREEN, line=PPT_BLUE, radius=True)
    add_textbox(slide, 4.77, 5.18, 3.8, 0.72, "Tracker\n6 drafts · 0 applied", font_size=13, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER, fill=PPT_AMBER, line=PPT_BLUE, radius=True)
    add_textbox(slide, 8.99, 5.18, 3.8, 0.72, "Recruiting Slack\n1 draft · 0 posted", font_size=13, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER, fill=PPT_LIGHT_BLUE, line=PPT_BLUE, radius=True)
    add_textbox(slide, 0.55, 6.12, 12.2, 0.55, "Slack: all eight names in source order + operational labels only; no IDs, emails, positions, raw result words or investigation reason", font_size=9.6, bold=True, color=PPT_NAVY, align=PP_ALIGN.CENTER, fill=PPT_RED, line=PPT_RED, radius=True)
    add_textbox(slide, 0.55, 6.84, 12.2, 0.33, "Oliver and Sanjay: email=0 · tracker update=0 · source status remains Awaiting Action", font_size=9.2, color=PPT_NAVY, align=PP_ALIGN.CENTER)
    presentation.save(path)
    normalize_ooxml(path)


def email_subject(item: Decision) -> str:
    if item.decision == "CLEARED_PROCEED":
        return f"Background Check Cleared - {item.source.candidate}"
    if item.decision == "FAILED_RESCIND":
        return f"HR Director Action Required - {item.source.candidate}"
    return f"Legal Notification Before Start - {item.source.candidate}"


def email_body(item: Decision) -> str:
    source = item.source
    common = [
        "DRAFT_NOT_SENT", "", f"Candidate: {source.candidate}", f"Candidate ID: {source.candidate_id}", f"Position: {source.position}", "",
    ]
    if item.decision == "CLEARED_PROCEED":
        detail = [
            "The background check is cleared. Onboarding may proceed under the Legal handling procedure.",
            "Proposed tracker status: Cleared - Proceed.",
        ]
    elif item.decision == "FAILED_RESCIND":
        detail = [
            "The background check requires the offer to be rescinded under the Legal handling procedure.",
            "Do not contact the candidate directly about the result.",
            "Proposed tracker status: Failed - Offer Rescinded.",
        ]
    else:
        detail = [
            "A vendor correction sets the effective result to Conditional. Onboarding may proceed, but Legal notification is required before the start date.",
            "The detailed vendor reason is intentionally omitted from this draft.",
            "Proposed tracker status: Conditional - Legal Notified.",
        ]
    return "\n".join(common + detail + ["", "This is an offline draft. No email has been sent and the source tracker has not been changed.", "- HR Operations"])


def write_mbox(path: Path, decisions: list[Decision]) -> None:
    chunks: list[bytes] = []
    order = 0
    for item in decisions:
        if item.email_draft != "YES":
            continue
        order += 1
        source = item.source
        message = EmailMessage(policy=SMTP)
        message["From"] = "hr-ops@company.example.com"
        message["To"] = item.notify_to
        message["Subject"] = email_subject(item)
        message["Date"] = "Wed, 15 Apr 2026 09:00:00 +0000"
        message["Message-ID"] = f"<background-5120-{source.source_row}@evocowork.local>"
        message["X-EvoCowork-Status"] = "DRAFT_NOT_SENT"
        message["X-EvoCowork-Source-Row"] = str(source.source_row)
        message["X-EvoCowork-Candidate-ID"] = source.candidate_id
        message["X-EvoCowork-Effective-Result"] = item.effective_result
        message["X-EvoCowork-Decision"] = item.decision
        message["X-EvoCowork-Proposed-Status"] = item.proposed_status
        message["X-EvoCowork-Order"] = str(order)
        message.set_content(email_body(item), charset="utf-8", cte="quoted-printable")
        stream = BytesIO()
        BytesGenerator(stream, policy=SMTP.clone(linesep="\n", max_line_length=78)).flatten(message, unixfrom=False)
        chunks.append(f"From MAILER-DAEMON Wed Apr 15 09:00:{order:02d} 2026\n".encode("ascii") + stream.getvalue().replace(b"\nFrom ", b"\n>From ") + b"\n")
    if order != 6:
        raise RuntimeError("背景调查邮件草稿必须恰好六封")
    path.write_bytes(b"".join(chunks))


def write_updates(path: Path, decisions: list[Decision]) -> None:
    drafts = []
    for item in decisions:
        if item.tracker_update != "YES":
            continue
        source = item.source
        drafts.append({
            "source_row": source.source_row,
            "candidate": source.candidate,
            "candidate_id": source.candidate_id,
            "source_result": source.source_result,
            "vendor_override": item.vendor_override,
            "effective_result": item.effective_result,
            "decision": item.decision,
            "source_status": source.source_status,
            "proposed_status": item.proposed_status,
            "notify_to": item.notify_to,
            "state": "DRAFT_NOT_APPLIED",
            "applied": False,
            "source_trace": source_trace(source.source_row).split(" | "),
        })
    payload = {
        "task_id": "automationbench-5120",
        "status": "CONTROLLED_DRAFT_ONLY",
        "frozen_time": FIXED_ISO_TIME,
        "precedence": ["VENDOR_CORRECTION", "LEGAL_RESULT_MAPPING", "AUDIENCE_AND_PRIVACY_CONTROL"],
        "draft_count": 6,
        "applied_count": 0,
        "source_modified": False,
        "drafts": drafts,
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_slack(path: Path, decisions: list[Decision]) -> None:
    entries = [{"source_row": item.source.source_row, "candidate": item.source.candidate, "operational_status": item.slack_label} for item in decisions]
    text = "#recruiting background-check operations update (offline draft):\n" + "\n".join(f"- {entry['candidate']} — {entry['operational_status']}" for entry in entries)
    payload = {
        "task_id": "automationbench-5120",
        "source_channel": {"id": "C_REC_5120", "name": "recruiting", "is_private": False},
        "privacy_mode": "NAMES_AND_OPERATIONAL_LABELS_ONLY",
        "draft_count": 1,
        "posted_count": 0,
        "posted": False,
        "drafts": [{
            "draft_id": "background-5120-recruiting-1",
            "status": "DRAFT_NOT_POSTED",
            "channel_id": "C_REC_5120",
            "channel_name": "recruiting",
            "entries": entries,
            "text": text,
        }],
        "excluded_fields": ["candidate_id", "email", "position", "raw_result", "investigation_reason"],
        "source_trace": ["Background_Check_Tracker.xlsx:Candidates rows 2-9", "Background_Check_Messages.mbox:msg_bgcheck_policy_5120", "Slack_Channels.json:C_REC_5120"],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_execution(path: Path) -> None:
    payload = {
        "task_id": "automationbench-5120",
        "status": "CONTROLLED_DRAFT_ONLY",
        "email": {"clear_drafts": 3, "failed_drafts": 2, "conditional_drafts": 1, "total_drafts": 6, "sent_count": 0, "state": "DRAFT_NOT_SENT"},
        "tracker_updates": {"clear_status_drafts": 3, "failed_status_drafts": 2, "conditional_status_drafts": 1, "total_drafts": 6, "applied_count": 0, "source_modified": False, "state": "DRAFT_NOT_APPLIED"},
        "pending_no_action": {"candidates": ["Oliver Strand", "Sanjay Gupta"], "source_status": "Awaiting Action", "email_draft_count": 0, "tracker_update_draft_count": 0},
        "slack": {"channel": "C_REC_5120 / recruiting", "draft_count": 1, "posted_count": 0, "state": "DRAFT_NOT_POSTED", "privacy_mode": "NAMES_AND_OPERATIONAL_LABELS_ONLY"},
        "source_tracker_modification_count": 0,
        "real_side_effect_count": 0,
        "source_trace": ["Background_Check_Tracker.xlsx:Candidates rows 2-9", "Background_Check_Messages.mbox:msg_bgcheck_correction_5120", "Background_Check_Messages.mbox:msg_bgcheck_policy_5120", "Workflow_Context.json:meta.current_time", "Slack_Channels.json:C_REC_5120"],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    records, messages, channels, frozen_time = read_inputs(args.input_dir)
    decisions = derive_decisions(records)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    write_workbook(args.output_dir / XLSX_NAME, records, decisions, messages, channels, frozen_time)
    write_brief(args.output_dir / DOCX_NAME, decisions)
    write_presentation(args.output_dir / PPTX_NAME, decisions)
    write_mbox(args.output_dir / MBOX_NAME, decisions)
    write_updates(args.output_dir / UPDATES_NAME, decisions)
    write_slack(args.output_dir / SLACK_NAME, decisions)
    write_execution(args.output_dir / EXECUTION_NAME)
    actual = {item.name for item in args.output_dir.iterdir() if item.is_file()}
    if actual != EXPECTED_OUTPUTS:
        raise RuntimeError(f"交付文件集合错误：{sorted(actual)}")


if __name__ == "__main__":
    main()
