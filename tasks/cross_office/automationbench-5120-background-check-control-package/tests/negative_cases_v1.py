#!/usr/bin/env python3
"""逐项注入 AutomationBench 5120 专属错误，验证 Judge 不会误接受坏答案。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
import zipfile
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path
from typing import Callable

from docx import Document
from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
XLSX = "Background_Check_Audit.xlsx"
DOCX = "Background_Check_Control_Brief.docx"
PPTX = "Background_Check_Review.pptx"
MBOX = "Background_Check_Notification_Drafts.mbox"
UPDATES = "Background_Check_Tracker_Update_Drafts.json"
SLACK = "Background_Check_Slack_Drafts.json"
EXECUTION = "Background_Check_Execution_Control.json"
FILES = (XLSX, DOCX, PPTX, MBOX, UPDATES, SLACK, EXECUTION)
Mutator = Callable[[Path], None]
Case = tuple[str, Mutator, str]


def load_judge():
    spec = importlib.util.spec_from_file_location("background_5120_judge", ROOT / "tests" / "judge.py")
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入本题 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    module.render = lambda path: ({".xlsx":10,".docx":3,".pptx":3}[path.suffix], "render-ok")
    return module


def rewrite_zip(path: Path, transform: Callable[[str, bytes], bytes | None]) -> None:
    temporary = path.with_suffix(path.suffix + ".negative")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = transform(info.filename, source.read(info.filename))
            if data is not None:
                target.writestr(info, data)
    temporary.replace(path)


def sheet_xml_path(archive: zipfile.ZipFile, sheet_name: str) -> str:
    workbook = etree.fromstring(archive.read("xl/workbook.xml"))
    relationships = etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    ns = {
        "m":"http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "r":"http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "pr":"http://schemas.openxmlformats.org/package/2006/relationships",
    }
    relationship_id = workbook.xpath(f'//m:sheet[@name="{sheet_name}"]/@r:id', namespaces=ns)[0]
    target = relationships.xpath(f'//pr:Relationship[@Id="{relationship_id}"]/@Target', namespaces=ns)[0]
    return target.lstrip("/") if target.startswith("/") else "xl/" + target.lstrip("/")


def edit_cell(path: Path, sheet_name: str, coordinate: str, *, value=None, remove_formula: bool = False) -> None:
    with zipfile.ZipFile(path, "r") as archive:
        target_sheet = sheet_xml_path(archive, sheet_name)
    changed = False
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

    def transform(name: str, data: bytes) -> bytes:
        nonlocal changed
        if name != target_sheet:
            return data
        root = etree.fromstring(data)
        cells = root.xpath(f'//m:c[@r="{coordinate}"]', namespaces={"m":namespace})
        if not cells:
            raise ValueError(f"{sheet_name}!{coordinate} 不存在")
        cell = cells[0]
        if remove_formula:
            formula_node = cell.find(f"{{{namespace}}}f")
            if formula_node is None:
                raise ValueError(f"{sheet_name}!{coordinate} 没有公式")
            cell.remove(formula_node)
        if value is not None:
            inline = cell.find(f"{{{namespace}}}is")
            if inline is not None:
                cell.remove(inline)
            node = cell.find(f"{{{namespace}}}v")
            if node is None:
                node = etree.SubElement(cell, f"{{{namespace}}}v")
            if isinstance(value, bool):
                node.text = "1" if value else "0"
                cell.set("t", "b")
            elif isinstance(value, (int, float)):
                node.text = str(value)
                cell.attrib.pop("t", None)
            else:
                node.text = str(value)
                cell.set("t", "str")
        changed = True
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    rewrite_zip(path, transform)
    if not changed:
        raise ValueError("目标单元格未修改")


def openpyxl_mutation(path: Path, mutator: Callable) -> None:
    workbook = load_workbook(path)
    mutator(workbook)
    workbook.save(path)


def mutate_json(path: Path, mutator: Callable[[dict], None]) -> None:
    payload = json.loads(path.read_text(encoding="utf-8"))
    mutator(payload)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def decode_message(message) -> str:
    if message.is_multipart():
        parts = []
        for part in message.walk():
            if part.get_content_type() == "text/plain":
                payload = part.get_payload(decode=True) or b""
                parts.append(payload.decode(part.get_content_charset() or "utf-8", errors="replace"))
        return "\n".join(parts).strip()
    payload = message.get_payload(decode=True) or b""
    return payload.decode(message.get_content_charset() or "utf-8", errors="replace").strip()


def mbox_states(path: Path) -> list[dict]:
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    return [{
        "from":str(message.get("From", "")),"to":str(message.get("To", "")),"cc":str(message.get("Cc", "")),"bcc":str(message.get("Bcc", "")),
        "subject":str(message.get("Subject", "")),"date":str(message.get("Date", "")),"message_id":str(message.get("Message-ID", "")),
        "status":str(message.get("X-EvoCowork-Status", "")),"source_row":str(message.get("X-EvoCowork-Source-Row", "")),
        "candidate_id":str(message.get("X-EvoCowork-Candidate-ID", "")),"effective_result":str(message.get("X-EvoCowork-Effective-Result", "")),
        "decision":str(message.get("X-EvoCowork-Decision", "")),"proposed_status":str(message.get("X-EvoCowork-Proposed-Status", "")),
        "order":str(message.get("X-EvoCowork-Order", "")),"body":decode_message(message),
    } for message in messages]


def write_mbox_states(path: Path, states: list[dict]) -> None:
    chunks = []
    headers = (
        ("from","From"),("to","To"),("cc","Cc"),("bcc","Bcc"),("subject","Subject"),("date","Date"),("message_id","Message-ID"),
        ("status","X-EvoCowork-Status"),("source_row","X-EvoCowork-Source-Row"),("candidate_id","X-EvoCowork-Candidate-ID"),
        ("effective_result","X-EvoCowork-Effective-Result"),("decision","X-EvoCowork-Decision"),("proposed_status","X-EvoCowork-Proposed-Status"),("order","X-EvoCowork-Order"),
    )
    for order, state in enumerate(states, 1):
        message = EmailMessage(policy=SMTP)
        for key, header in headers:
            if state.get(key):
                message[header] = state[key]
        message.set_content(state["body"].rstrip() + "\n", charset="utf-8")
        buffer = BytesIO()
        BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
        chunks.extend([f"From MAILER-DAEMON Wed Apr 15 09:00:{order:02d} 2026\n".encode("ascii"),buffer.getvalue().replace(b"\r\n", b"\n"),b"\n"])
    path.write_bytes(b"".join(chunks))


def mutate_message(path: Path, index: int, mutator: Callable[[dict], None]) -> None:
    states = mbox_states(path)
    mutator(states[index])
    write_mbox_states(path, states)


def replace_docx_text(path: Path, old: str, new: str, *, row_token: str | None = None) -> None:
    document = Document(path)
    changed = False
    if row_token is None:
        for paragraph in document.paragraphs:
            if old in paragraph.text:
                paragraph.text = paragraph.text.replace(old, new, 1)
                changed = True
                break
    if not changed:
        for table in document.tables:
            for row in table.rows:
                if row_token and row_token not in " | ".join(cell.text for cell in row.cells):
                    continue
                for cell in row.cells:
                    if old in cell.text:
                        cell.text = cell.text.replace(old, new, 1)
                        changed = True
                        break
                if changed:
                    break
            if changed:
                break
    if not changed:
        raise ValueError(f"Word 找不到 {old!r}")
    document.save(path)


def replace_pptx_text(path: Path, old: str, new: str, *, row_token: str | None = None) -> None:
    presentation = Presentation(path)
    changed = False
    for slide in presentation.slides:
        for shape in slide.shapes:
            if row_token is None and getattr(shape, "has_text_frame", False) and old in shape.text:
                shape.text = shape.text.replace(old, new, 1)
                changed = True
                break
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    if row_token and row_token not in " | ".join(cell.text for cell in row.cells):
                        continue
                    for cell in row.cells:
                        if old in cell.text:
                            cell.text = cell.text.replace(old, new, 1)
                            changed = True
                            break
                    if changed:
                        break
            if changed:
                break
        if changed:
            break
    if not changed:
        raise ValueError(f"PPT 找不到 {old!r}")
    presentation.save(path)


def drop_ppt_slide(path: Path) -> None:
    presentation = Presentation(path)
    presentation.slides._sldIdLst.remove(presentation.slides._sldIdLst[-1])
    presentation.save(path)


def build_cases() -> list[Case]:
    cases: list[Case] = []
    for name in FILES:
        cases.append((f"missing-{name}", lambda root, n=name: (root / n).unlink(), "R001"))
    cases.append(("extra-file", lambda root: (root / "unexpected.txt").write_text("unexpected", encoding="utf-8"), "R001"))
    corrupt_criterion = {XLSX:"R001",DOCX:"R001",PPTX:"R001",MBOX:"R010",UPDATES:"R011",SLACK:"R012",EXECUTION:"R013"}
    for name in FILES:
        cases.append((f"corrupt-{name}", lambda root, n=name: (root / n).write_bytes(b"not-a-valid-artifact"), corrupt_criterion[name]))

    source_wrong = {"A":99,"B":"Wrong Person","C":"BGC-9999","D":"Wrong Manager","E":"wrong@company.example.com","F":"Wrong Position","G":"Wrong Result","H":"Wrong Status","I":"Wrong Note"}
    for row in range(2, 10):
        for column, value in source_wrong.items():
            cases.append((f"source-{column}{row}", lambda root, c=f"{column}{row}", v=value: edit_cell(root / XLSX, "Source_Candidates", c, value=v), "R002"))

    for row in (2,3):
        for column, value in (("A","wrong-id"),("B","wrong@company.example.com"),("C","wrong@company.example.com"),("D","Wrong subject"),("E","Thu, 16 Apr 2026 09:00:00 +0000"),("F","Wrong body"),("G","Wrong authority"),("H","Wrong privacy")):
            cases.append((f"message-{column}{row}", lambda root, c=f"{column}{row}", v=value: edit_cell(root / XLSX, "Message_Source", c, value=v), "R003"))
    for row in range(2, 18):
        for column, value in (("A","WRONG-ID"),("D","Wrong fact"),("E","Wrong use"),("F","Wrong privacy"),("G","wrong-source")):
            cases.append((f"evidence-{column}{row}", lambda root, c=f"{column}{row}", v=value: edit_cell(root / XLSX, "Evidence_Register", c, value=v), "R003"))

    for row in range(2, 10):
        for column in range(1, 25):
            letter = chr(64 + column)
            cases.append((f"formula-{letter}{row}", lambda root, c=f"{letter}{row}": edit_cell(root / XLSX, "Candidate_Decisions", c, remove_formula=True), "R004"))

    people = {2:"R005",3:"R006",4:"R007",5:"R008",6:"R005",7:"R006",8:"R005",9:"R008"}
    decision_wrong = {"J":"Wrong Override","K":"Wrong Effective","L":"WRONG","M":"WRONG","N":"WRONG","O":"WRONG","P":"WRONG_DECISION","Q":"Wrong Status","R":"wrong@company.example.com","S":"WRONG","T":"WRONG","U":"WRONG_LABEL","V":"Wrong privacy","W":"Wrong authority","X":"wrong-source"}
    for row, criterion in people.items():
        for column, value in decision_wrong.items():
            cases.append((f"decision-{column}{row}", lambda root, c=f"{column}{row}", v=value: edit_cell(root / XLSX, "Candidate_Decisions", c, value=v), criterion))

    for row in range(2, 5):
        for column, value in (("C","Wrong Person"),("D","BGC-9999"),("E","Wrong Manager"),("F","wrong@company.example.com"),("G","WRONG"),("H","Wrong Status"),("I","EXECUTED"),("J","wrong-source")):
            cases.append((f"clear-{column}{row}", lambda root, c=f"{column}{row}", v=value: edit_cell(root / XLSX, "Clear_Queue", c, value=v), "R005"))
    for row in range(2, 4):
        for column, value in (("C","Wrong Person"),("D","BGC-9999"),("F","YES"),("G","wrong@company.example.com"),("H","ALLOWED"),("I","WRONG"),("J","Wrong Status"),("K","WRONG_LABEL"),("L","EXECUTED"),("M","wrong-source")):
            cases.append((f"failed-{column}{row}", lambda root, c=f"{column}{row}", v=value: edit_cell(root / XLSX, "Failed_Control", c, value=v), "R006"))
    for column, value in (("B","Wrong Naomi"),("C","BGC-9999"),("D","Failed"),("E","Clear"),("F","Clear"),("H","YES"),("I","lisa.wang@company.example.com"),("J","CLEARED_PROCEED"),("K","Cleared - Proceed"),("L","DISCLOSED"),("M","ONBOARDING_MAY_PROCEED"),("N","EXECUTED"),("O","wrong-source")):
        cases.append((f"conditional-{column}", lambda root, c=f"{column}2", v=value: edit_cell(root / XLSX, "Conditional_Control", c, value=v), "R007"))
    for row in range(2, 4):
        for column, value in (("B","Wrong Pending"),("D","Clear"),("E","Cleared - Proceed"),("F","Clear"),("G","CLEARED_PROCEED"),("H","Cleared - Proceed"),("I","YES"),("J","YES"),("K","ONBOARDING_MAY_PROCEED"),("L","EXECUTED"),("M","wrong-source")):
            cases.append((f"pending-{column}{row}", lambda root, c=f"{column}{row}", v=value: edit_cell(root / XLSX, "Pending_Control", c, value=v), "R008"))
    for row in range(2, 10):
        for column, value in (("C","Wrong Person"),("D","Wrong Effective"),("E","WRONG"),("F","wrong@company.example.com"),("G","SENT"),("H","APPLIED"),("I","WRONG_LABEL"),("J","POSTED"),("K","Wrong privacy"),("L","wrong-source")):
            cases.append((f"communication-{column}{row}", lambda root, c=f"{column}{row}", v=value: edit_cell(root / XLSX, "Communication_Register", c, value=v), "R009"))
    for row in range(2, 15):
        criterion = "R013" if row in (8,10,12,13,14) else "R009"
        cases.append((f"summary-B{row}", lambda root, c=f"B{row}": edit_cell(root / XLSX, "Control_Summary", c, value=99), criterion))
    cases.extend([
        ("xlsx-table-removed", lambda root: openpyxl_mutation(root / XLSX, lambda wb: wb["Source_Candidates"].tables.clear()), "R017"),
        ("xlsx-chart-removed", lambda root: openpyxl_mutation(root / XLSX, lambda wb: setattr(wb["Control_Summary"], "_charts", [])), "R017"),
        ("xlsx-cf-removed", lambda root: openpyxl_mutation(root / XLSX, lambda wb: setattr(wb["Candidate_Decisions"], "conditional_formatting", type(wb["Candidate_Decisions"].conditional_formatting)())), "R017"),
    ])

    cases.extend([
        ("mbox-empty", lambda root: write_mbox_states(root / MBOX, []), "R010"),
        ("mbox-drop-first", lambda root: write_mbox_states(root / MBOX, mbox_states(root / MBOX)[1:]), "R010"),
        ("mbox-drop-last", lambda root: write_mbox_states(root / MBOX, mbox_states(root / MBOX)[:-1]), "R010"),
        ("mbox-duplicate", lambda root: write_mbox_states(root / MBOX, mbox_states(root / MBOX) + [mbox_states(root / MBOX)[-1]]), "R010"),
        ("mbox-reverse", lambda root: write_mbox_states(root / MBOX, list(reversed(mbox_states(root / MBOX)))), "R010"),
    ])
    action_criteria = ["R005","R006","R007","R005","R006","R005"]
    for index, criterion in enumerate(action_criteria):
        for field, value in (("to","wrong@company.example.com"),("cc","cc@company.example.com"),("bcc","bcc@company.example.com"),("subject","Wrong subject"),("date","Thu, 16 Apr 2026 09:00:00 +0000"),("message_id","<wrong@evocowork.local>"),("status","SENT"),("source_row","99"),("candidate_id","BGC-9999"),("effective_result","Wrong"),("decision","WRONG"),("proposed_status","Wrong Status"),("order","99")):
            cases.append((f"mbox-{index}-{field}", lambda root, i=index, f=field, v=value: mutate_message(root / MBOX, i, lambda state: state.update({f:v})), criterion))
    body_mutations = [
        (0,"Liam Chen","Wrong Person","R005"),(0,"Cleared - Proceed","Failed - Offer Rescinded","R005"),
        (1,"requires the offer to be rescinded","requires onboarding to proceed","R006"),(1,"Do not contact the candidate directly","Contact the candidate directly","R006"),
        (2,"Conditional - Legal Notified","Cleared - Proceed","R007"),(2,"detailed vendor reason is intentionally omitted","pending court disposition","R007"),
        (3,"Petra Johansson","Wrong Person","R005"),(4,"Quinn Davis","Wrong Person","R006"),(5,"Rosa Martinez","Wrong Person","R005"),
        (0,"DRAFT_NOT_SENT","SENT","R005"),
    ]
    for index, old, new, criterion in body_mutations:
        def mutate_body(root, i=index, o=old, n=new):
            def change(state):
                if o not in state["body"]:
                    raise ValueError(f"邮件正文找不到 {o!r}")
                state["body"] = state["body"].replace(o, n, 1)
            mutate_message(root / MBOX, i, change)
        cases.append((f"mbox-body-{index}-{old[:12]}", mutate_body, criterion))

    for key, value in (("task_id","wrong"),("status","APPLIED"),("frozen_time","2026-04-16T09:00:00Z"),("draft_count",5),("applied_count",1),("source_modified",True)):
        cases.append((f"updates-{key}", lambda root, k=key, v=value: mutate_json(root / UPDATES, lambda payload: payload.update({k:v})), "R011"))
    cases.extend([
        ("updates-precedence", lambda root: mutate_json(root / UPDATES, lambda payload: payload["precedence"].reverse()), "R011"),
        ("updates-drop-first", lambda root: mutate_json(root / UPDATES, lambda payload: payload["drafts"].pop(0)), "R011"),
        ("updates-drop-last", lambda root: mutate_json(root / UPDATES, lambda payload: payload["drafts"].pop()), "R011"),
        ("updates-reverse", lambda root: mutate_json(root / UPDATES, lambda payload: payload["drafts"].reverse()), "R011"),
        ("updates-add-oliver", lambda root: mutate_json(root / UPDATES, lambda payload: payload["drafts"].append({"candidate":"Oliver Strand"})), "R008"),
        ("updates-add-sanjay", lambda root: mutate_json(root / UPDATES, lambda payload: payload["drafts"].append({"candidate":"Sanjay Gupta"})), "R008"),
    ])
    draft_wrong = {"source_row":99,"candidate":"Wrong Person","candidate_id":"BGC-9999","source_result":"Wrong","vendor_override":"Wrong","effective_result":"Wrong","decision":"WRONG","source_status":"Wrong","proposed_status":"Wrong","notify_to":"wrong@company.example.com","state":"APPLIED","applied":True,"source_trace":[]}
    for index, criterion in enumerate(action_criteria):
        for key, value in draft_wrong.items():
            cases.append((f"updates-{index}-{key}", lambda root, i=index, k=key, v=value: mutate_json(root / UPDATES, lambda payload: payload["drafts"][i].update({k:v})), criterion))

    for key, value in (("task_id","wrong"),("privacy_mode","FULL_DETAILS"),("draft_count",2),("posted_count",1),("posted",True),("excluded_fields",[]),("source_trace",[])):
        cases.append((f"slack-{key}", lambda root, k=key, v=value: mutate_json(root / SLACK, lambda payload: payload.update({k:v})), "R012"))
    cases.extend([
        ("slack-channel-id", lambda root: mutate_json(root / SLACK, lambda payload: payload["source_channel"].update({"id":"C_OTHER"})), "R012"),
        ("slack-channel-private", lambda root: mutate_json(root / SLACK, lambda payload: payload["source_channel"].update({"is_private":True})), "R012"),
        ("slack-drop-entry", lambda root: mutate_json(root / SLACK, lambda payload: payload["drafts"][0]["entries"].pop()), "R012"),
        ("slack-reverse-entry", lambda root: mutate_json(root / SLACK, lambda payload: payload["drafts"][0]["entries"].reverse()), "R012"),
        ("slack-wrong-label", lambda root: mutate_json(root / SLACK, lambda payload: payload["drafts"][0]["entries"][2].update({"operational_status":"ONBOARDING_MAY_PROCEED"})), "R012"),
        ("slack-leak-id", lambda root: mutate_json(root / SLACK, lambda payload: payload["drafts"][0].update({"text":payload["drafts"][0]["text"] + " BGC-4412"})), "R012"),
        ("slack-leak-email", lambda root: mutate_json(root / SLACK, lambda payload: payload["drafts"][0].update({"text":payload["drafts"][0]["text"] + " lisa.wang@company.example.com"})), "R012"),
        ("slack-leak-position", lambda root: mutate_json(root / SLACK, lambda payload: payload["drafts"][0].update({"text":payload["drafts"][0]["text"] + " Account Executive"})), "R012"),
        ("slack-leak-result", lambda root: mutate_json(root / SLACK, lambda payload: payload["drafts"][0].update({"text":payload["drafts"][0]["text"] + " Conditional"})), "R012"),
        ("slack-leak-reason", lambda root: mutate_json(root / SLACK, lambda payload: payload["drafts"][0].update({"text":payload["drafts"][0]["text"] + " pending court disposition"})), "R012"),
    ])

    execution_mutations = [
        ("task-id", lambda payload: payload.update({"task_id":"wrong"})),("status", lambda payload: payload.update({"status":"EXECUTED"})),("side-effect", lambda payload: payload.update({"real_side_effect_count":1})),
        ("email-clear", lambda payload: payload["email"].update({"clear_drafts":2})),("email-failed", lambda payload: payload["email"].update({"failed_drafts":3})),("email-conditional", lambda payload: payload["email"].update({"conditional_drafts":0})),
        ("email-total", lambda payload: payload["email"].update({"total_drafts":5})),("email-sent", lambda payload: payload["email"].update({"sent_count":1,"state":"SENT"})),
        ("tracker-clear", lambda payload: payload["tracker_updates"].update({"clear_status_drafts":2})),("tracker-failed", lambda payload: payload["tracker_updates"].update({"failed_status_drafts":3})),
        ("tracker-conditional", lambda payload: payload["tracker_updates"].update({"conditional_status_drafts":0})),("tracker-total", lambda payload: payload["tracker_updates"].update({"total_drafts":5})),
        ("tracker-applied", lambda payload: payload["tracker_updates"].update({"applied_count":1,"state":"APPLIED"})),("tracker-modified", lambda payload: payload["tracker_updates"].update({"source_modified":True})),
        ("pending-list", lambda payload: payload["pending_no_action"].update({"candidates":[]})),("pending-status", lambda payload: payload["pending_no_action"].update({"source_status":"Cleared - Proceed"})),
        ("pending-email", lambda payload: payload["pending_no_action"].update({"email_draft_count":1})),("pending-update", lambda payload: payload["pending_no_action"].update({"tracker_update_draft_count":1})),
        ("slack-draft", lambda payload: payload["slack"].update({"draft_count":2})),("slack-post", lambda payload: payload["slack"].update({"posted_count":1,"state":"POSTED"})),
        ("source-modification", lambda payload: payload.update({"source_tracker_modification_count":1})),("trace", lambda payload: payload.update({"source_trace":[]})),
    ]
    for name, mutator in execution_mutations:
        criterion = "R008" if name.startswith("pending") else "R013"
        cases.append((f"execution-{name}", lambda root, m=mutator: mutate_json(root / EXECUTION, m), criterion))

    word_cases = [
        ("clear-count","3","4","Clear","R015"),("failed-count","2","3","Failed","R015"),("conditional-count","1","0","Conditional","R015"),("pending-count","2","1","Pending","R015"),
        ("liam-route","derek.huang@company.example.com","wrong@company.example.com","Liam Chen","R005"),("maya-route","hr-director@company.example.com","mira.patel@company.example.com","Maya Okafor","R006"),
        ("naomi-override","Conditional","Clear","BGC-4412","R007"),("naomi-route","legal@company.example.com","lisa.wang@company.example.com","BGC-4412","R007"),
        ("oliver-action","PENDING_NO_ACTION","CLEARED_PROCEED","Oliver Strand","R008"),("petra-route","mira.patel@company.example.com","wrong@company.example.com","Petra Johansson","R005"),
        ("quinn-route","hr-director@company.example.com","lisa.wang@company.example.com","Quinn Davis","R006"),("rosa-route","derek.huang@company.example.com","wrong@company.example.com","Rosa Martinez","R005"),
        ("sanjay-action","PENDING_NO_ACTION","CLEARED_PROCEED","Sanjay Gupta","R008"),("email-state","DRAFT_NOT_SENT","SENT","Emails","R013"),
        ("tracker-state","DRAFT_NOT_APPLIED","APPLIED","Tracker updates","R013"),("slack-state","DRAFT_NOT_POSTED","POSTED","Recruiting Slack","R013"),("side-effect","OFFLINE_ONLY","EXECUTED","Real side effects","R013"),
    ]
    for name, old, new, row_token, criterion in word_cases:
        cases.append((f"word-{name}", lambda root, o=old, n=new, token=row_token: replace_docx_text(root / DOCX, o, n, row_token=token), criterion))
    cases.extend([
        ("word-naomi-manager", lambda root: replace_docx_text(root / DOCX, "Lisa Wang is not notified", "Lisa Wang is notified"), "R007"),
        ("word-pending-action", lambda root: replace_docx_text(root / DOCX, "generate neither email nor tracker-update drafts", "generate email and tracker updates"), "R008"),
        ("word-sensitive-leak", lambda root: replace_docx_text(root / DOCX, "Her detailed vendor reason remains restricted", "Her pending court disposition is shared"), "R007"),
    ])

    ppt_cases = [
        ("summary-clear","3","4","R014"),("summary-failed","2","3","R014"),
        ("naomi-precedence","tracker Clear -> vendor correction Conditional -> Legal only","tracker Clear -> manager","R007"),("naomi-manager","Lisa Wang is not notified","Lisa Wang is notified","R007"),
        ("email-count","6 drafts · 0 sent","5 drafts · 1 sent","R013"),("tracker-count","6 drafts · 0 applied","5 drafts · 1 applied","R013"),
        ("slack-count","1 draft · 0 posted","1 draft · 1 posted","R013"),("pending-action","Oliver and Sanjay: email=0","Oliver and Sanjay: email=1","R008"),
    ]
    for name, old, new, criterion in ppt_cases:
        row_token = "Liam; Petra; Rosa" if name == "summary-clear" else ("Maya; Quinn" if name == "summary-failed" else None)
        cases.append((f"ppt-{name}", lambda root, o=old, n=new, token=row_token: replace_pptx_text(root / PPTX, o, n, row_token=token), criterion))
    row_cases = [
        ("liam-route","Liam Chen","derek.huang@company.example.com","wrong@company.example.com","R005"),("maya-route","Maya Okafor","hr-director@company.example.com","mira.patel@company.example.com","R006"),
        ("naomi-override","Naomi Yoshida","Conditional","Clear","R007"),("naomi-route","Naomi Yoshida","legal@company.example.com","lisa.wang@company.example.com","R007"),
        ("oliver-decision","Oliver Strand","PENDING_NO_ACTION","CLEARED_PROCEED","R008"),("petra-route","Petra Johansson","mira.patel@company.example.com","wrong@company.example.com","R005"),
        ("quinn-route","Quinn Davis","hr-director@company.example.com","lisa.wang@company.example.com","R006"),("rosa-route","Rosa Martinez","derek.huang@company.example.com","wrong@company.example.com","R005"),
        ("sanjay-decision","Sanjay Gupta","PENDING_NO_ACTION","CLEARED_PROCEED","R008"),
    ]
    for name, token, old, new, criterion in row_cases:
        cases.append((f"ppt-{name}", lambda root, t=token, o=old, n=new: replace_pptx_text(root / PPTX, o, n, row_token=t), criterion))
    cases.extend([
        ("ppt-sensitive-leak", lambda root: replace_pptx_text(root / PPTX, "Sensitive vendor reason stays restricted", "pending court disposition shared"), "R007"),
        ("ppt-drop-slide", lambda root: drop_ppt_slide(root / PPTX), "R014"),
    ])
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--json-out", type=Path)
    args = parser.parse_args()
    judge = load_judge()
    cases = build_cases()
    results, failures = [], []
    for name, mutator, expected_failure in cases:
        with tempfile.TemporaryDirectory(prefix=f"negative-5120-{name}-") as temporary:
            artifact_root = Path(temporary) / "artifacts"
            shutil.copytree(GOLDEN, artifact_root)
            try:
                mutator(artifact_root)
                result = judge.evaluate(artifact_root)
                detected = not result.get("criteria", {}).get(expected_failure, {}).get("passed", True)
                rejected = not result.get("passed", False) and float(result.get("reward", 1.0)) < 0.8
                results.append({"case":name,"expected_failure":expected_failure,"detected":detected,"rejected":rejected,"reward":result.get("reward"),"hurdle_failures":result.get("hurdle_failures",[])})
                if not detected or not rejected:
                    failures.append(f"{name}: expected={expected_failure}, detected={detected}, reward={result.get('reward')}")
            except Exception as exc:
                failures.append(f"{name}: mutation/evaluation error: {type(exc).__name__}: {exc}")
    payload = {
        "task_id":"cowork-evo/automationbench-5120-background-check-control-package",
        "case_count":len(results),"declared_case_count":len(cases),"all_rejected":not failures and len(results) == len(cases),
        "max_reward":max((float(item["reward"]) for item in results), default=None),"failures":failures,"results":results,
    }
    if args.json_out:
        args.json_out.parent.mkdir(parents=True, exist_ok=True)
        args.json_out.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    if failures or len(results) != len(cases):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
