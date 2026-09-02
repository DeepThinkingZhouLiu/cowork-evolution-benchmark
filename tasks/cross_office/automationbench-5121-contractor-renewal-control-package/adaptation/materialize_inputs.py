#!/usr/bin/env python3
"""只把 AutomationBench 5121 的公开初始状态静态化为本题输入。

脚本逐项断言八份合同、采购政策、冻结时间、HR Slack 快照和 13 条上游
assertions。它不是通用制题器，不会创造合同工、日期、PIP、费率或审批人。
"""

from __future__ import annotations

import hashlib
import importlib.util
import json
import shutil
import sys
import types
import zipfile
from datetime import date, datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
UPSTREAM = Path(__file__).with_name("upstream-hr-tasks.py")
SNAPSHOT = Path(__file__).with_name("upstream-task-5121.json")
LICENSE = Path(__file__).with_name("LICENSE-AutomationBench-MIT")
DATA_DIR = ROOT / "data" / "input_files"
ENV_DIR = ROOT / "environment" / "input_files"
FIXED_ZIP_TIME = (2026, 4, 15, 9, 0, 0)
SOURCE_SHA256 = "c173a36a03f43096bd122577584aa32f9df379d8760aa46c43be04c85d101779"
LICENSE_SHA256 = "c6f9d87aff6d627d5cd78cb09364dafed7d6f1c4bb7a91183d57377157e96ba4"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_task() -> dict:
    if sha256(UPSTREAM) != SOURCE_SHA256 or sha256(LICENSE) != LICENSE_SHA256:
        raise ValueError("冻结的 AutomationBench 源文件或 LICENSE 哈希错误")
    datasets = types.ModuleType("datasets")
    datasets.Dataset = object
    sys.modules["datasets"] = datasets
    for name in ("automationbench", "automationbench.domains", "automationbench.domains.hr"):
        sys.modules[name] = types.ModuleType(name)
    noise = types.ModuleType("automationbench.domains.hr._noise")
    noise.apply_noise = lambda value: value
    sys.modules[noise.__name__] = noise
    spec = importlib.util.spec_from_file_location("automationbench_hr_tasks_frozen_5121", UPSTREAM)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入冻结的 AutomationBench HR 源文件")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    task = module.get_task_5121()
    if task.get("example_id") != 5121 or task.get("task") != "hr.contractor_renewal":
        raise ValueError("冻结函数不再是 AutomationBench 5121")
    return task


def assert_shape(task: dict) -> None:
    prompt = task["prompt"][1]["content"]
    assert all(token in prompt for token in ("Renew all the contractor agreements", "Send renewal notices", "standard 5% annual rate increase", "contractor management guidelines"))
    state = task["info"]["initial_state"]
    assert state["meta"] == {"schema_version":"0.1.0","current_time":"2026-04-15T09:00:00Z"}
    messages = state["gmail"]["messages"]
    assert len(messages) == 1
    policy = messages[0]
    assert policy["id"] == "msg_contractor_policy_5121"
    assert policy["from_"] == "procurement@company.example.com"
    assert policy["subject"] == "Contractor Agreement Renewal Policy"
    assert all(token in policy["body_plain"] for token in (
        "expiring within the next 60 days", "active Performance Improvement Plan", "Pending VP Approval",
        "Agreements expiring beyond 60 days: No action", "Expired - Not Renewed",
        "HR Ops may NOT change contractor rates", "including standard cost-of-living increases",
        "Procurement (procurement@company.example.com) and Finance approval", "existing terms only",
    ))
    spreadsheet = state["google_sheets"]["spreadsheets"][0]
    assert spreadsheet["id"] == "ss_contractors_5121" and spreadsheet["title"] == "Contractor Agreements"
    worksheet = spreadsheet["worksheets"][0]
    assert worksheet["id"] == "ws_agreements_5121" and worksheet["title"] == "Active Agreements"
    rows = worksheet["rows"]
    assert [item["row_id"] for item in rows] == list(range(2, 10))
    assert [item["cells"]["Contractor"] for item in rows] == ["Ana Petrova","Brian Osei","Carmen Ruiz","Darius Kim","Eva Lindström","Felix Braun","Gloria Chen","Hugo Santos"]
    assert [item["cells"]["Agreement End"] for item in rows] == ["2026-05-15","2026-06-30","2026-04-30","2026-05-01","2026-08-15","2026-04-10","2026-06-10","2026-05-20"]
    assert [item["cells"]["Notes"] for item in rows] == ["","","On PIP since March 2026","","","","On PIP since February 2026",""]
    assert state["slack"] == {"channels":[{"id":"C_HR_5121","name":"hr-general","is_private":False}],"messages":[]}
    assertions = task["info"]["assertions"]
    assert len(assertions) == 13
    serialized = json.dumps(assertions)
    assert all(token in serialized for token in (
        "ana.petrova@freelance.example.com", "brian.osei@freelance.example.com",
        "Carmen Ruiz", "darius.kim@freelance.example.com",
        "eva.lindstrom@freelance.example.com", "Expired - Not Renewed",
        "Gloria Chen", "Pending VP Approval", "5%",
    ))


def normalize_xlsx(path: Path) -> None:
    temporary = path.with_suffix(".normalized.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    temporary.replace(path)


def style_sheet(sheet, widths: dict[str, float]) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = sheet.dimensions


def write_tracker(state: dict, target: Path) -> None:
    source = state["google_sheets"]["spreadsheets"][0]
    worksheet = source["worksheets"][0]
    workbook = Workbook()
    workbook.creator = "AutomationBench 5121 public fixture"
    workbook.title = source["title"]
    workbook.subject = "Static Office representation of contractor agreement tracker"
    workbook.description = "No contractor, email, department, end date, status or PIP note was added beyond AutomationBench 5121."
    sheet = workbook.active
    sheet.title = worksheet["title"]
    headers = ["Source Row","Contractor","Contractor Email","Manager","Manager Email","Department","Dept VP Email","Agreement End","Status","Notes"]
    sheet.append(headers)
    for item in worksheet["rows"]:
        cells = item["cells"]
        end = date.fromisoformat(cells["Agreement End"])
        sheet.append([item["row_id"],cells["Contractor"],cells["Contractor Email"],cells["Manager"],cells["Manager Email"],cells["Department"],cells["Dept VP Email"],end,cells["Status"],cells["Notes"]])
    for row in range(2, 10):
        sheet.cell(row, 8).number_format = "yyyy-mm-dd"
        sheet.row_dimensions[row].height = 28
    table = Table(displayName="ContractorAgreementSourceInput", ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    style_sheet(sheet, {"A":12,"B":24,"C":40,"D":22,"E":38,"F":20,"G":38,"H":18,"I":16,"J":34})
    workbook.save(target)
    normalize_xlsx(target)


def message_bytes(record: dict) -> bytes:
    message = EmailMessage(policy=SMTP)
    message["From"] = record["from_"]
    message["To"] = ", ".join(record["to"])
    message["Subject"] = record["subject"]
    parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
    message["Date"] = parsed.strftime("%a, %d %b %Y %H:%M:%S +0000")
    message["Message-ID"] = f"<{record['id']}@automationbench.invalid>"
    message["X-AutomationBench-Message-ID"] = record["id"]
    message["X-AutomationBench-Thread-ID"] = record["thread_id"]
    message["X-AutomationBench-Labels"] = ",".join(record.get("label_ids", []))
    message.set_content(record["body_plain"], charset="utf-8")
    buffer = BytesIO()
    BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
    return buffer.getvalue()


def write_mbox(messages: list[dict], target: Path) -> None:
    chunks: list[bytes] = []
    for record in messages:
        parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
        chunks.append(f"From {record['from_']} {parsed.strftime('%a %b %d %H:%M:%S %Y')}\n".encode("ascii"))
        chunks.append(message_bytes(record).replace(b"\r\n", b"\n"))
        chunks.append(b"\n")
    target.write_bytes(b"".join(chunks))


def main() -> None:
    task = load_task()
    assert_shape(task)
    state = task["info"]["initial_state"]
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    ENV_DIR.mkdir(parents=True, exist_ok=True)
    SNAPSHOT.write_text(json.dumps(task, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_tracker(state, DATA_DIR / "Contractor_Agreement_Tracker.xlsx")
    write_mbox(state["gmail"]["messages"], DATA_DIR / "Contractor_Renewal_Policy.mbox")
    (DATA_DIR / "Slack_Channels.json").write_text(json.dumps(state["slack"], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (DATA_DIR / "Workflow_Context.json").write_text(json.dumps({"meta":state["meta"]}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    for source in sorted(DATA_DIR.iterdir()):
        if source.is_file():
            shutil.copy2(source, ENV_DIR / source.name)


if __name__ == "__main__":
    main()
