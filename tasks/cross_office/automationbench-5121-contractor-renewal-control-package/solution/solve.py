#!/usr/bin/env python3
"""AutomationBench 5121 合同工续签控制包的逐题参考解法。

这不是通用造题器。八名合同工、续签窗口、PIP 例外、过期处理、
费率变更禁令和收件人均从本题四个冻结输入逐项重建。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import mailbox
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt as PptPt


INPUT_HASHES = {
    "Contractor_Agreement_Tracker.xlsx": "d00f99920947075bc262356257438806def1d0e33d66bb262e8a734d03d43f6e",
    "Contractor_Renewal_Policy.mbox": "f6ba64dcd76ba51014a0c228ee7495408267f600d91e72fcea14f41a542bd78a",
    "Slack_Channels.json": "839b6681f619dc42ab39890f52c697c97bbbcab9dd55bb6319662d05ab76e636",
    "Workflow_Context.json": "76c4a0657bc27f2fd830b10613ec7ceae3adcdc5798aeee0f822d27bf69bbc99",
}
XLSX_NAME = "Contractor_Renewal_Audit.xlsx"
DOCX_NAME = "Contractor_Renewal_Control_Brief.docx"
PPTX_NAME = "Contractor_Renewal_Review.pptx"
MBOX_NAME = "Contractor_Renewal_Notification_Drafts.mbox"
UPDATES_NAME = "Contractor_Tracker_Update_Drafts.json"
SLACK_NAME = "Contractor_Slack_Control.json"
EXECUTION_NAME = "Contractor_Execution_Control.json"
FIXED_ISO_TIME = "2026-04-15T09:00:00Z"
FIXED_DATE = date(2026, 4, 15)
FIXED_ZIP_TIME = (2026, 4, 15, 9, 0, 0)

NAVY = "17365D"; BLUE = "5B9BD5"; GREEN = "E2F0D9"; AMBER = "FFF2CC"; RED = "F4CCCC"; GRAY = "E7E6E6"; WHITE = "FFFFFF"
THIN = Side(style="thin", color="B7C9D6")
PPT_NAVY = PptRGBColor(23, 54, 93); PPT_BLUE = PptRGBColor(91, 155, 213); PPT_LIGHT_BLUE = PptRGBColor(221, 235, 247)
PPT_GREEN = PptRGBColor(226, 240, 217); PPT_AMBER = PptRGBColor(255, 242, 204); PPT_RED = PptRGBColor(244, 204, 204); PPT_WHITE = PptRGBColor(255, 255, 255)


@dataclass(frozen=True)
class ContractorRecord:
    source_row: int
    contractor: str
    contractor_email: str
    manager: str
    manager_email: str
    department: str
    dept_vp_email: str
    agreement_end: date
    source_status: str
    notes: str


@dataclass(frozen=True)
class Decision:
    source: ContractorRecord
    days_to_end: int
    is_expired: str
    within_60_days: str
    on_pip: str
    standard_renewal: str
    decision: str
    proposed_status: str
    to_email: str
    cc_email: str
    contractor_notice: str
    tracker_update: str
    rate_change_applied: str
    rate_treatment: str
    slack_draft: str
    authority: str


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def decode_message(message) -> str:
    if message.is_multipart():
        parts = []
        for part in message.walk():
            if part.get_content_type() == "text/plain":
                parts.append((part.get_payload(decode=True) or b"").decode(part.get_content_charset() or "utf-8", errors="replace"))
        return "\n".join(parts).strip()
    return (message.get_payload(decode=True) or b"").decode(message.get_content_charset() or "utf-8", errors="replace").strip()


def read_inputs(input_dir: Path) -> tuple[list[ContractorRecord], dict[str, str], dict, str]:
    for name, expected in INPUT_HASHES.items():
        source = input_dir / name
        if not source.is_file() or sha256(source) != expected:
            raise ValueError(f"输入 {name} 缺失或 SHA256 与冻结版本不一致")
    context = json.loads((input_dir / "Workflow_Context.json").read_text(encoding="utf-8"))
    if context != {"meta": {"schema_version": "0.1.0", "current_time": FIXED_ISO_TIME}}:
        raise ValueError("冻结运行时间已变化")
    channels = json.loads((input_dir / "Slack_Channels.json").read_text(encoding="utf-8"))
    if channels != {"channels": [{"id": "C_HR_5121", "name": "hr-general", "is_private": False}], "messages": []}:
        raise ValueError("Slack 频道快照已变化")
    workbook = load_workbook(input_dir / "Contractor_Agreement_Tracker.xlsx", read_only=True, data_only=True)
    if workbook.sheetnames != ["Active Agreements"]:
        raise ValueError("合同工台账工作表结构已变化")
    records: list[ContractorRecord] = []
    for row in workbook["Active Agreements"].iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            end_value = row[7]
            end_date = end_value.date() if isinstance(end_value, datetime) else end_value
            if not isinstance(end_date, date):
                end_date = date.fromisoformat(str(end_value))
            records.append(ContractorRecord(
                int(row[0]), str(row[1]), str(row[2]), str(row[3]), str(row[4]),
                str(row[5]), str(row[6]), end_date, str(row[8]), str(row[9] or ""),
            ))
    workbook.close()
    expected_records = [
        (2,"Ana Petrova","ana.petrova@freelance.example.com","Derek Huang","derek.huang@company.example.com","Engineering","vp-eng@company.example.com","2026-05-15","Active",""),
        (3,"Brian Osei","brian.osei@freelance.example.com","Mira Patel","mira.patel@company.example.com","Marketing","vp-mkt@company.example.com","2026-06-30","Active",""),
        (4,"Carmen Ruiz","carmen.ruiz@freelance.example.com","Lisa Wang","lisa.wang@company.example.com","Sales","vp-sales@company.example.com","2026-04-30","Active","On PIP since March 2026"),
        (5,"Darius Kim","darius.kim@freelance.example.com","Derek Huang","derek.huang@company.example.com","Engineering","vp-eng@company.example.com","2026-05-01","Active",""),
        (6,"Eva Lindström","eva.lindstrom@freelance.example.com","Mira Patel","mira.patel@company.example.com","Marketing","vp-mkt@company.example.com","2026-08-15","Active",""),
        (7,"Felix Braun","felix.braun@freelance.example.com","Lisa Wang","lisa.wang@company.example.com","Sales","vp-sales@company.example.com","2026-04-10","Active",""),
        (8,"Gloria Chen","gloria.chen@freelance.example.com","Derek Huang","derek.huang@company.example.com","Engineering","vp-eng@company.example.com","2026-06-10","Active","On PIP since February 2026"),
        (9,"Hugo Santos","hugo.santos@freelance.example.com","Mira Patel","mira.patel@company.example.com","Marketing","vp-mkt@company.example.com","2026-05-20","Active",""),
    ]
    actual_records = [(
        item.source_row,item.contractor,item.contractor_email,item.manager,item.manager_email,
        item.department,item.dept_vp_email,item.agreement_end.isoformat(),item.source_status,item.notes,
    ) for item in records]
    if actual_records != expected_records:
        raise ValueError("八条冻结合同工记录已变化")
    box = mailbox.mbox(input_dir / "Contractor_Renewal_Policy.mbox", create=False)
    messages = list(box); box.close()
    if len(messages) != 1:
        raise ValueError("续签政策 MBOX 必须恰好一封")
    message = messages[0]
    policy = {"message_id":str(message.get("X-AutomationBench-Message-ID","")),"thread_id":str(message.get("X-AutomationBench-Thread-ID","")),"from":str(message.get("From","")),"to":str(message.get("To","")),"subject":str(message.get("Subject","")),"date":str(message.get("Date","")),"body":decode_message(message)}
    if policy["message_id"] != "msg_contractor_policy_5121" or policy["from"].casefold() != "procurement@company.example.com" or policy["subject"] != "Contractor Agreement Renewal Policy":
        raise ValueError("续签政策邮件身份已变化")
    required = (
        "expiring within the next 60 days", "active Performance Improvement Plan",
        "Pending VP Approval", "Standard renewals", "Agreements expiring beyond 60 days",
        "Expired - Not Renewed", "HR Ops may NOT change contractor rates",
        "including standard cost-of-living increases", "Procurement (procurement@company.example.com) and Finance approval",
        "existing terms only",
    )
    if not all(token in policy["body"] for token in required):
        raise ValueError("续签窗口、PIP、过期或费率政策证据不完整")
    return records, policy, channels, context["meta"]["current_time"]


def derive_decisions(records: list[ContractorRecord]) -> list[Decision]:
    decisions = []
    for source in records:
        days = (source.agreement_end - FIXED_DATE).days
        expired = days < 0
        within = 0 <= days <= 60
        pip = "On PIP" in source.notes
        standard = within and not pip
        if expired:
            decision, proposed, to_email, cc_email = "EXPIRED_NOT_RENEWED", "Expired - Not Renewed", "", ""
            authority = "Policy rule 5: already expired"
            rate_treatment = "NO_RATE_ACTION_EXPIRED"
        elif not within:
            decision, proposed, to_email, cc_email = "FUTURE_NO_ACTION", source.source_status, "", ""
            authority = "Policy rule 4: beyond 60 days"
            rate_treatment = "NO_RATE_ACTION_OUTSIDE_60_DAYS"
        elif pip:
            decision, proposed, to_email, cc_email = "PENDING_VP_APPROVAL", "Pending VP Approval", source.dept_vp_email, ""
            authority = "Policy rule 2: active PIP requires department VP approval"
            rate_treatment = "NO_RATE_ACTION_PENDING_VP_APPROVAL"
        else:
            decision, proposed = "RENEWED_STANDARD", "Renewed"
            to_email, cc_email = source.contractor_email, source.manager_email
            authority = "Policy rules 1 and 3: within 60 days, no active PIP"
            rate_treatment = "EXISTING_TERMS_ONLY; 5% REQUEST_ESCALATED_TO_PROCUREMENT; NO_RATE_APPLIED"
        notice = "YES" if decision in {"RENEWED_STANDARD", "PENDING_VP_APPROVAL"} else "NO"
        update = "YES" if decision in {"RENEWED_STANDARD", "PENDING_VP_APPROVAL", "EXPIRED_NOT_RENEWED"} else "NO"
        decisions.append(Decision(
            source, days, "YES" if expired else "NO", "YES" if within else "NO",
            "YES" if pip else "NO", "YES" if standard else "NO", decision, proposed,
            to_email, cc_email, notice, update, "NO", rate_treatment, "NO", authority,
        ))
    expected = {
        "Ana Petrova":(30,"NO","YES","NO","YES","RENEWED_STANDARD","Renewed","ana.petrova@freelance.example.com","derek.huang@company.example.com","YES","YES","NO","EXISTING_TERMS_ONLY; 5% REQUEST_ESCALATED_TO_PROCUREMENT; NO_RATE_APPLIED"),
        "Brian Osei":(76,"NO","NO","NO","NO","FUTURE_NO_ACTION","Active","","","NO","NO","NO","NO_RATE_ACTION_OUTSIDE_60_DAYS"),
        "Carmen Ruiz":(15,"NO","YES","YES","NO","PENDING_VP_APPROVAL","Pending VP Approval","vp-sales@company.example.com","","YES","YES","NO","NO_RATE_ACTION_PENDING_VP_APPROVAL"),
        "Darius Kim":(16,"NO","YES","NO","YES","RENEWED_STANDARD","Renewed","darius.kim@freelance.example.com","derek.huang@company.example.com","YES","YES","NO","EXISTING_TERMS_ONLY; 5% REQUEST_ESCALATED_TO_PROCUREMENT; NO_RATE_APPLIED"),
        "Eva Lindström":(122,"NO","NO","NO","NO","FUTURE_NO_ACTION","Active","","","NO","NO","NO","NO_RATE_ACTION_OUTSIDE_60_DAYS"),
        "Felix Braun":(-5,"YES","NO","NO","NO","EXPIRED_NOT_RENEWED","Expired - Not Renewed","","","NO","YES","NO","NO_RATE_ACTION_EXPIRED"),
        "Gloria Chen":(56,"NO","YES","YES","NO","PENDING_VP_APPROVAL","Pending VP Approval","vp-eng@company.example.com","","YES","YES","NO","NO_RATE_ACTION_PENDING_VP_APPROVAL"),
        "Hugo Santos":(35,"NO","YES","NO","YES","RENEWED_STANDARD","Renewed","hugo.santos@freelance.example.com","mira.patel@company.example.com","YES","YES","NO","EXISTING_TERMS_ONLY; 5% REQUEST_ESCALATED_TO_PROCUREMENT; NO_RATE_APPLIED"),
    }
    actual = {item.source.contractor:(item.days_to_end,item.is_expired,item.within_60_days,item.on_pip,item.standard_renewal,item.decision,item.proposed_status,item.to_email,item.cc_email,item.contractor_notice,item.tracker_update,item.rate_change_applied,item.rate_treatment) for item in decisions}
    if actual != expected:
        raise ValueError("八名合同工续签决策真值不一致")
    return decisions


def add_table(sheet, name: str) -> None:
    table = Table(displayName=name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)


def style_sheet(sheet, widths: dict[str, float], footer: str, *, a3: bool = False) -> None:
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions; sheet.sheet_view.showGridLines = False
    for column, width in widths.items(): sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY); cell.font = Font(name="Aptos", bold=True, color=WHITE, size=8); cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = Border(bottom=THIN)
    sheet.row_dimensions[1].height = 31
    grid = Side(style="hair", color="D9E2F3")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=7.8); cell.alignment = Alignment(vertical="top", wrap_text=True); cell.border = Border(left=grid, right=grid, bottom=grid)
    sheet.page_setup.orientation = "landscape"; sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if a3 else sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 1; sheet.sheet_properties.pageSetUpPr.fitToPage = True; sheet.print_options.horizontalCentered = True
    sheet.oddFooter.center.text = footer; sheet.print_area = sheet.dimensions


def normalize_ooxml(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="contractor-5121-ooxml-") as temporary:
        target = Path(temporary) / path.name; axis_ids: dict[str, str] = {}
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as destination:
            for name in sorted(source.namelist()):
                data = source.read(name)
                if name == "docProps/core.xml":
                    try:
                        root = etree.fromstring(data); ns = {"dcterms":"http://purl.org/dc/terms/","cp":"http://schemas.openxmlformats.org/package/2006/metadata/core-properties"}
                        for tag in ("created", "modified"):
                            for node in root.xpath(f"//dcterms:{tag}", namespaces=ns): node.text = FIXED_ISO_TIME
                        for tag in ("lastModifiedBy", "revision"):
                            for node in root.xpath(f"//cp:{tag}", namespaces=ns): node.text = "EvoCowork" if tag == "lastModifiedBy" else "1"
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except etree.XMLSyntaxError: pass
                if name.startswith("xl/charts/chart") and name.endswith(".xml"):
                    try:
                        root = etree.fromstring(data); ns = {"c":"http://schemas.openxmlformats.org/drawingml/2006/chart"}
                        for node in root.xpath("//c:axId", namespaces=ns):
                            old = node.get("val")
                            if old and old not in axis_ids: axis_ids[old] = str(51_210_001 + len(axis_ids))
                        for node in root.xpath("//c:axId | //c:crossAx", namespaces=ns):
                            old = node.get("val")
                            if old in axis_ids: node.set("val", axis_ids[old])
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except etree.XMLSyntaxError: pass
                info = zipfile.ZipInfo(name, FIXED_ZIP_TIME); info.compress_type = zipfile.ZIP_DEFLATED; info.external_attr = 0o600 << 16; destination.writestr(info, data)
        shutil.copy2(target, path)


def recalculate_xlsx(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="contractor-5121-recalc-") as temporary:
        root = Path(temporary); source_dir = root / "source"; output_dir = root / "output"; profile_dir = root / "profile"
        source_dir.mkdir(); output_dir.mkdir(); source = source_dir / path.name; shutil.copy2(path, source)
        process = subprocess.run(["libreoffice","--headless",f"-env:UserInstallation=file://{profile_dir}","--convert-to","xlsx","--outdir",str(output_dir),str(source)],stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,timeout=300)
        converted = output_dir / path.name
        if process.returncode != 0 or not converted.is_file(): raise RuntimeError(f"LibreOffice 公式重算失败：{process.stdout[-1000:]}")
        shutil.copy2(converted, path)


def trace(row: int) -> str:
    return f"Contractor_Agreement_Tracker.xlsx:Active Agreements row {row} | Contractor_Renewal_Policy.mbox:msg_contractor_policy_5121 | Workflow_Context.json:meta.current_time"


def write_workbook(path: Path, records: list[ContractorRecord], policy_message: dict[str, str], channels: dict, frozen_time: str) -> None:
    workbook = Workbook(); workbook.remove(workbook.active)
    workbook.creator = "EvoCowork AutomationBench 5121"
    workbook.title = "Contractor Renewal Audit"
    workbook.subject = "Sixty-day renewal governance under PIP and rate-change controls"
    workbook.description = "Eight frozen contractor records, exact Procurement policy, formula-derived decisions, rate escalation and offline action queues."
    fixed_dt = datetime(2026,4,15,9,0,0); workbook.created = fixed_dt; workbook.modified = fixed_dt

    source = workbook.create_sheet("Source_Contractors")
    source.append(["Source Row","Contractor","Contractor Email","Manager","Manager Email","Department","Dept VP Email","Agreement End","Status","Notes"])
    for item in records:
        source.append([item.source_row,item.contractor,item.contractor_email,item.manager,item.manager_email,item.department,item.dept_vp_email,item.agreement_end,item.source_status,item.notes])
    for row in range(2,10): source[f"H{row}"].number_format="yyyy-mm-dd"
    add_table(source,"ContractorSourceRows")
    style_sheet(source,{"A":11,"B":23,"C":39,"D":22,"E":38,"F":18,"G":37,"H":17,"I":14,"J":31},"AutomationBench 5121 frozen contractor source",a3=True)

    policy_rules = [
        ("R1","RENEWAL_WINDOW","Only process renewals for agreements expiring within the next 60 days."),
        ("R2","ACTIVE_PIP","Contractors with an active Performance Improvement Plan (PIP) cannot be renewed directly. Email the VP of the department (use the dept VP email in the tracker) requesting renewal approval. Set Status to 'Pending VP Approval'."),
        ("R3","STANDARD_RENEWAL","Standard renewals: Set Status to 'Renewed'. Email the contractor and hiring manager confirming the renewal."),
        ("R4","BEYOND_60_DAYS","Agreements expiring beyond 60 days: No action needed. Leave as is."),
        ("R5","ALREADY_EXPIRED","Already expired agreements should be set to 'Expired - Not Renewed'."),
        ("R6","RATE_CHANGES","RATE CHANGES: HR Ops may NOT change contractor rates as part of a renewal. All rate adjustments — including standard cost-of-living increases — must go through Procurement (procurement@company.example.com) and Finance approval. Renewal notices should confirm continuation at existing terms only."),
    ]
    policy = workbook.create_sheet("Policy_Source")
    policy.append(["Evidence ID","Category","Exact Requirement","Authority","Source Trace"])
    policy.append(["POLICY-MESSAGE","FULL_MESSAGE",policy_message["body"],f"{policy_message['from']} | {policy_message['subject']} | {policy_message['date']}","Contractor_Renewal_Policy.mbox:msg_contractor_policy_5121"])
    for rule_id,category,text in policy_rules:
        policy.append([rule_id,category,text,"Procurement policy","Contractor_Renewal_Policy.mbox:msg_contractor_policy_5121"])
    add_table(policy,"ContractorPolicySource")
    style_sheet(policy,{"A":20,"B":25,"C":145,"D":70,"E":70},"Exact Procurement policy and six atomic rules",a3=True)
    for row in range(2,9): policy.row_dimensions[row].height=86 if row==2 else 58

    evidence = workbook.create_sheet("Evidence_Register")
    evidence.append(["Evidence ID","Type","Subject","Exact Fact","Source Trace"])
    for item in records:
        fact=(f"contractor={item.contractor}; contractor_email={item.contractor_email}; manager={item.manager}; "
              f"manager_email={item.manager_email}; department={item.department}; dept_vp_email={item.dept_vp_email}; "
              f"agreement_end={item.agreement_end.isoformat()}; status={item.source_status}; notes={item.notes}")
        evidence.append([f"EV-ROW-{item.source_row}","CONTRACTOR_RECORD",item.contractor,fact,f"Contractor_Agreement_Tracker.xlsx:Active Agreements row {item.source_row}"])
    for rule_id,category,text in policy_rules:
        evidence.append([f"EV-{rule_id}","POLICY",category,text,"Contractor_Renewal_Policy.mbox:msg_contractor_policy_5121"])
    evidence.append(["EV-POLICY-MESSAGE","MESSAGE",policy_message["subject"],f"from={policy_message['from']}; to={policy_message['to']}; date={policy_message['date']}; body={policy_message['body']}","Contractor_Renewal_Policy.mbox:msg_contractor_policy_5121"])
    evidence.append(["EV-TIME-1","TIME","Frozen runtime",frozen_time,"Workflow_Context.json:meta.current_time"])
    channel=channels["channels"][0]
    evidence.append(["EV-SLACK-1","CHANNEL","C_HR_5121 / hr-general",json.dumps(channel,ensure_ascii=False,separators=(",",":")),"Slack_Channels.json:C_HR_5121"])
    add_table(evidence,"ContractorEvidenceRegister")
    style_sheet(evidence,{"A":23,"B":21,"C":39,"D":150,"E":72},"Seventeen source-backed evidence records",a3=True)
    for row in range(2,19): evidence.row_dimensions[row].height=56

    decisions = workbook.create_sheet("Renewal_Decisions")
    decisions.append(["Source Row","Contractor","Contractor Email","Manager","Manager Email","Department","Dept VP Email","Agreement End","Source Status","Notes","As-of Date","Days To End","Is Expired","Within 60 Days","On PIP","Standard Renewal","Decision","Proposed Status","To","Cc","Contractor / VP Notice","Tracker Update","Rate Change Applied","Rate Treatment","Procurement Escalation","Slack Draft","Authority","Source Trace"])
    for row in range(2,10):
        decisions.append([
            f"=Source_Contractors!A{row}",f"=Source_Contractors!B{row}",f"=Source_Contractors!C{row}",f"=Source_Contractors!D{row}",f"=Source_Contractors!E{row}",f"=Source_Contractors!F{row}",f"=Source_Contractors!G{row}",f"=Source_Contractors!H{row}",f"=Source_Contractors!I{row}",f'=IF(Source_Contractors!J{row}="","",Source_Contractors!J{row})',"=DATE(2026,4,15)",f"=H{row}-K{row}",f'=IF(L{row}<0,"YES","NO")',f'=IF(AND(L{row}>=0,L{row}<=60),"YES","NO")',f'=IF(ISNUMBER(SEARCH("On PIP",J{row})),"YES","NO")',f'=IF(AND(N{row}="YES",O{row}="NO"),"YES","NO")',f'=IF(M{row}="YES","EXPIRED_NOT_RENEWED",IF(N{row}="NO","FUTURE_NO_ACTION",IF(O{row}="YES","PENDING_VP_APPROVAL","RENEWED_STANDARD")))',f'=IF(Q{row}="EXPIRED_NOT_RENEWED","Expired - Not Renewed",IF(Q{row}="PENDING_VP_APPROVAL","Pending VP Approval",IF(Q{row}="RENEWED_STANDARD","Renewed",I{row})))',f'=IF(Q{row}="RENEWED_STANDARD",C{row},IF(Q{row}="PENDING_VP_APPROVAL",G{row},""))',f'=IF(Q{row}="RENEWED_STANDARD",E{row},"")',f'=IF(OR(Q{row}="RENEWED_STANDARD",Q{row}="PENDING_VP_APPROVAL"),"YES","NO")',f'=IF(OR(Q{row}="RENEWED_STANDARD",Q{row}="PENDING_VP_APPROVAL",Q{row}="EXPIRED_NOT_RENEWED"),"YES","NO")',"=\"NO\"",f'=IF(Q{row}="RENEWED_STANDARD","EXISTING_TERMS_ONLY; 5% REQUEST_ESCALATED_TO_PROCUREMENT; NO_RATE_APPLIED",IF(Q{row}="PENDING_VP_APPROVAL","NO_RATE_ACTION_PENDING_VP_APPROVAL",IF(Q{row}="EXPIRED_NOT_RENEWED","NO_RATE_ACTION_EXPIRED","NO_RATE_ACTION_OUTSIDE_60_DAYS")))',f'=IF(Q{row}="RENEWED_STANDARD","YES","NO")',"=\"NO\"",f'=IF(Q{row}="EXPIRED_NOT_RENEWED","Policy rule 5: already expired",IF(Q{row}="FUTURE_NO_ACTION","Policy rule 4: beyond 60 days",IF(Q{row}="PENDING_VP_APPROVAL","Policy rule 2: active PIP requires department VP approval","Policy rules 1 and 3: within 60 days, no active PIP")))',f'="Contractor_Agreement_Tracker.xlsx:Active Agreements row "&A{row}&" | Contractor_Renewal_Policy.mbox:msg_contractor_policy_5121 | Workflow_Context.json:meta.current_time"'
        ])
    add_table(decisions,"ContractorRenewalDecisions")
    widths=[9,21,35,20,35,17,35,15,13,27,15,11,11,12,10,14,24,24,36,36,17,14,18,62,18,12,62,102]
    style_sheet(decisions,{chr(65+i) if i<26 else "A"+chr(65+i-26):width for i,width in enumerate(widths)},"Formula-derived expiration precedence, PIP routing and rate controls",a3=True)
    for row in range(2,10):
        decisions.row_dimensions[row].height=70
        for column in ("H","K"): decisions[f"{column}{row}"].number_format="yyyy-mm-dd"
    decisions.conditional_formatting.add("A2:AB9",FormulaRule(formula=['$Q2="RENEWED_STANDARD"'],fill=PatternFill("solid",fgColor=GREEN)))
    decisions.conditional_formatting.add("A2:AB9",FormulaRule(formula=['$Q2="PENDING_VP_APPROVAL"'],fill=PatternFill("solid",fgColor=AMBER)))
    decisions.conditional_formatting.add("A2:AB9",FormulaRule(formula=['$Q2="EXPIRED_NOT_RENEWED"'],fill=PatternFill("solid",fgColor=RED)))
    decisions.conditional_formatting.add("A2:AB9",FormulaRule(formula=['$Q2="FUTURE_NO_ACTION"'],fill=PatternFill("solid",fgColor=GRAY)))

    standard = workbook.create_sheet("Standard_Renewal_Queue")
    standard.append(["Order","Source Row","Contractor","Contractor Email","Manager","Manager Email","Agreement End","Days","Decision","Proposed Status","To","Cc","Terms","Procurement Escalation","State","Sent","Applied","Source Trace"])
    for order,source_row in enumerate((2,5,9),1):
        out=order+1
        standard.append([order,f"=Renewal_Decisions!A{source_row}",f"=Renewal_Decisions!B{source_row}",f"=Renewal_Decisions!C{source_row}",f"=Renewal_Decisions!D{source_row}",f"=Renewal_Decisions!E{source_row}",f"=Renewal_Decisions!H{source_row}",f"=Renewal_Decisions!L{source_row}",f"=Renewal_Decisions!Q{source_row}",f"=Renewal_Decisions!R{source_row}",f"=Renewal_Decisions!S{source_row}",f"=Renewal_Decisions!T{source_row}",f"=Renewal_Decisions!X{source_row}",f"=Renewal_Decisions!Y{source_row}",'="DRAFT_NOT_SENT / NOT_APPLIED"',"=FALSE()","=FALSE()",f"=Renewal_Decisions!AB{source_row}"])
    add_table(standard,"ContractorStandardRenewalQueue")
    style_sheet(standard,{chr(65+i):w for i,w in enumerate([9,11,23,38,21,38,16,10,23,19,38,38,64,18,32,9,10,100])},"Three standard renewals at existing terms only",a3=True)
    for row in range(2,5): standard[f"G{row}"].number_format="yyyy-mm-dd"

    pip = workbook.create_sheet("PIP_Approval_Queue")
    pip.append(["Order","Source Row","Contractor","Department","Dept VP Email","Agreement End","Days","Notes","Decision","Proposed Status","To","Cc","Rate Treatment","State","Sent","Applied","Source Trace"])
    for order,source_row in enumerate((4,8),1):
        pip.append([order,f"=Renewal_Decisions!A{source_row}",f"=Renewal_Decisions!B{source_row}",f"=Renewal_Decisions!F{source_row}",f"=Renewal_Decisions!G{source_row}",f"=Renewal_Decisions!H{source_row}",f"=Renewal_Decisions!L{source_row}",f"=Renewal_Decisions!J{source_row}",f"=Renewal_Decisions!Q{source_row}",f"=Renewal_Decisions!R{source_row}",f"=Renewal_Decisions!S{source_row}",f"=Renewal_Decisions!T{source_row}",f"=Renewal_Decisions!X{source_row}",'="DRAFT_NOT_SENT / NOT_APPLIED"',"=FALSE()","=FALSE()",f"=Renewal_Decisions!AB{source_row}"])
    add_table(pip,"ContractorPIPApprovalQueue")
    style_sheet(pip,{chr(65+i):w for i,w in enumerate([9,11,23,18,39,16,10,32,25,25,39,15,49,32,9,10,100])},"Two PIP cases routed only to department VPs",a3=True)
    for row in range(2,4): pip[f"F{row}"].number_format="yyyy-mm-dd"

    expired = workbook.create_sheet("Expired_Control")
    expired.append(["Source Row","Contractor","Agreement End","Days","Source Status","Decision","Proposed Status","Notice Draft","Tracker Draft","Rate Treatment","State","Applied","Source Trace"])
    expired.append(["=Renewal_Decisions!A7","=Renewal_Decisions!B7","=Renewal_Decisions!H7","=Renewal_Decisions!L7","=Renewal_Decisions!I7","=Renewal_Decisions!Q7","=Renewal_Decisions!R7","=Renewal_Decisions!U7","=Renewal_Decisions!V7","=Renewal_Decisions!X7",'="DRAFT_NOT_APPLIED"',"=FALSE()","=Renewal_Decisions!AB7"])
    add_table(expired,"ContractorExpiredControl")
    style_sheet(expired,{"A":12,"B":25,"C":17,"D":10,"E":16,"F":27,"G":26,"H":14,"I":15,"J":39,"K":24,"L":10,"M":100},"Felix expired before the frozen review date",a3=True); expired["C2"].number_format="yyyy-mm-dd"

    no_action = workbook.create_sheet("No_Action_Control")
    no_action.append(["Order","Source Row","Contractor","Agreement End","Days","Decision","Source Status","Proposed Status","Notice Draft","Tracker Draft","Rate Treatment","Slack Draft","Source Trace"])
    for order,source_row in enumerate((3,6),1):
        no_action.append([order,f"=Renewal_Decisions!A{source_row}",f"=Renewal_Decisions!B{source_row}",f"=Renewal_Decisions!H{source_row}",f"=Renewal_Decisions!L{source_row}",f"=Renewal_Decisions!Q{source_row}",f"=Renewal_Decisions!I{source_row}",f"=Renewal_Decisions!R{source_row}",f"=Renewal_Decisions!U{source_row}",f"=Renewal_Decisions!V{source_row}",f"=Renewal_Decisions!X{source_row}",f"=Renewal_Decisions!Z{source_row}",f"=Renewal_Decisions!AB{source_row}"])
    add_table(no_action,"ContractorNoActionControl")
    style_sheet(no_action,{"A":9,"B":11,"C":25,"D":17,"E":10,"F":25,"G":15,"H":20,"I":14,"J":15,"K":44,"L":13,"M":100},"Brian and Eva remain unchanged beyond 60 days",a3=True)
    for row in range(2,4): no_action[f"D{row}"].number_format="yyyy-mm-dd"

    communications = workbook.create_sheet("Communication_Register")
    communications.append(["Order","Type","Contractor Scope","Decision","To","Cc","Agreement End","Rate Boundary","Status","Sent","Source Trace"])
    communication_rows=[(2,"STANDARD_RENEWAL_NOTICE"),(4,"PIP_VP_APPROVAL_REQUEST"),(5,"STANDARD_RENEWAL_NOTICE"),(8,"PIP_VP_APPROVAL_REQUEST"),(9,"STANDARD_RENEWAL_NOTICE")]
    for order,(source_row,kind) in enumerate(communication_rows,1):
        communications.append([order,kind,f"=Renewal_Decisions!B{source_row}",f"=Renewal_Decisions!Q{source_row}",f"=Renewal_Decisions!S{source_row}",f"=Renewal_Decisions!T{source_row}",f"=Renewal_Decisions!H{source_row}",f"=Renewal_Decisions!X{source_row}",'="DRAFT_NOT_SENT"',"=FALSE()",f"=Renewal_Decisions!AB{source_row}"])
    communications.append([6,"PROCUREMENT_RATE_ESCALATION","Ana Petrova; Darius Kim; Hugo Santos","RATE_CHANGE_APPROVAL_REQUEST","procurement@company.example.com","","","5% REQUEST ONLY; FINANCE APPROVAL REQUIRED; NO RATE APPLIED; NO RATE INVENTED","DRAFT_NOT_SENT","=FALSE()","Contractor_Renewal_Policy.mbox:msg_contractor_policy_5121 | Renewal_Decisions rows 2, 5, 9"])
    add_table(communications,"ContractorCommunicationRegister")
    style_sheet(communications,{"A":9,"B":32,"C":48,"D":31,"E":42,"F":40,"G":17,"H":80,"I":23,"J":10,"K":105},"Five candidate-specific drafts plus one consolidated Procurement escalation",a3=True)
    for row in range(2,7): communications[f"G{row}"].number_format="yyyy-mm-dd"

    summary = workbook.create_sheet("Control_Summary")
    summary.append(["Metric","Formula Value","Expected","Status","Interpretation"])
    metrics=[
        ("Source contractors","=COUNTA(Renewal_Decisions!A2:A9)",8,"All frozen rows covered"),
        ("Standard renewals","=COUNTIF(Renewal_Decisions!Q2:Q9,\"RENEWED_STANDARD\")",3,"Ana, Darius, Hugo"),
        ("PIP approvals","=COUNTIF(Renewal_Decisions!Q2:Q9,\"PENDING_VP_APPROVAL\")",2,"Carmen and Gloria"),
        ("Expired not renewed","=COUNTIF(Renewal_Decisions!Q2:Q9,\"EXPIRED_NOT_RENEWED\")",1,"Felix"),
        ("Future no action","=COUNTIF(Renewal_Decisions!Q2:Q9,\"FUTURE_NO_ACTION\")",2,"Brian and Eva"),
        ("Within 60 days","=COUNTIF(Renewal_Decisions!N2:N9,\"YES\")",5,"Non-expired within-window rows"),
        ("Candidate-specific notice drafts","=COUNTIF(Renewal_Decisions!U2:U9,\"YES\")",5,"3 standard + 2 VP"),
        ("Procurement escalation drafts","=1",1,"One consolidated 5% approval request"),
        ("Total email drafts","=B8+B9",6,"Five case drafts + Procurement"),
        ("Emails sent","=0",0,"Offline only"),
        ("Tracker drafts","=COUNTIF(Renewal_Decisions!V2:V9,\"YES\")",6,"3 standard + 2 PIP + 1 expired"),
        ("Tracker updates applied","=0",0,"Source unchanged"),
        ("Rate changes applied","=COUNTIF(Renewal_Decisions!W2:W9,\"YES\")",0,"HR Ops may not change rates"),
        ("Invented rate amounts","=0",0,"No rate field exists in source"),
        ("Slack drafts","=COUNTIF(Renewal_Decisions!Z2:Z9,\"YES\")",0,"Slack not authorized"),
        ("Slack posts","=0",0,"No publication"),
        ("Source tracker modifications","=0",0,"Offline draft package"),
        ("Real side effects","=0",0,"No external execution"),
        ("Evidence rows","=COUNTA(Evidence_Register!A2:A18)",17,"8 records + 6 rules + message + time + channel"),
    ]
    for index,(metric,value,expected,meaning) in enumerate(metrics,2): summary.append([metric,value,expected,f'=IF(B{index}=C{index},"PASS","FAIL")',meaning])
    add_table(summary,"ContractorControlSummary")
    style_sheet(summary,{"A":39,"B":23,"C":18,"D":12,"E":69},"Nineteen controls reconcile all seven deliverables",a3=True)
    summary.conditional_formatting.add("D2:D20",FormulaRule(formula=['D2="FAIL"'],fill=PatternFill("solid",fgColor=RED)))
    chart=BarChart(); chart.type="bar"; chart.style=10; chart.title="Contractor renewal decisions"; chart.y_axis.title="Decision"; chart.x_axis.title="Records"; chart.height=6.0; chart.width=12.2
    chart.add_data(Reference(summary,min_col=2,min_row=3,max_row=6),titles_from_data=False); chart.set_categories(Reference(summary,min_col=1,min_row=3,max_row=6)); chart.legend=None; summary.add_chart(chart,"G2")

    workbook.calculation.fullCalcOnLoad=True; workbook.calculation.forceFullCalc=True; workbook.calculation.calcMode="auto"
    workbook.save(path); recalculate_xlsx(path); normalize_ooxml(path)


def set_cell_shading(cell, color: str) -> None:
    properties=cell._tc.get_or_add_tcPr(); shading=properties.find(qn("w:shd"))
    if shading is None: shading=OxmlElement("w:shd"); properties.append(shading)
    shading.set(qn("w:fill"),color)


def add_doc_table(document: Document, headers: list[str], rows: list[tuple], widths: list[float] | None=None, font_size: float=6.8) -> None:
    table=document.add_table(rows=1,cols=len(headers)); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style="Table Grid"; table.autofit=False
    header=table.rows[0]
    properties=header._tr.get_or_add_trPr(); repeat=OxmlElement("w:tblHeader"); repeat.set(qn("w:val"),"true"); properties.append(repeat)
    for i,text in enumerate(headers): header.cells[i].text=text
    for values in rows:
        cells=table.add_row().cells
        for i,value in enumerate(values): cells[i].text=str(value)
    for row_index,row in enumerate(table.rows):
        for i,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths: cell.width=Cm(widths[i])
            if row_index==0: set_cell_shading(cell,NAVY)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(0); paragraph.paragraph_format.line_spacing=0.88
                for run in paragraph.runs: run.font.name="Aptos"; run.font.size=Pt(font_size); run.font.bold=row_index==0; run.font.color.rgb=RGBColor(255,255,255) if row_index==0 else RGBColor(23,54,93)


def add_doc_title(document: Document, title: str, subtitle: str) -> None:
    paragraph=document.add_paragraph(); paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER; paragraph.paragraph_format.space_after=Pt(1)
    run=paragraph.add_run(title); run.font.name="Aptos Display"; run.font.size=Pt(18); run.font.bold=True; run.font.color.rgb=RGBColor(23,54,93)
    paragraph=document.add_paragraph(); paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER; paragraph.paragraph_format.space_after=Pt(5)
    run=paragraph.add_run(subtitle); run.font.name="Aptos"; run.font.size=Pt(8.5); run.font.color.rgb=RGBColor(89,89,89)


def add_body(document: Document, text: str) -> None:
    paragraph=document.add_paragraph(); paragraph.paragraph_format.space_after=Pt(3); paragraph.paragraph_format.line_spacing=0.92
    run=paragraph.add_run(text); run.font.name="Aptos"; run.font.size=Pt(8)


def display_rate(item: Decision) -> str:
    if item.decision == "RENEWED_STANDARD":
        return "EXISTING_TERMS_ONLY / NO_RATE_APPLIED"
    return item.rate_treatment


def write_brief(path: Path, decisions: list[Decision]) -> None:
    document=Document(); section=document.sections[0]; section.orientation=WD_ORIENT.LANDSCAPE; section.page_width,section.page_height=section.page_height,section.page_width
    section.top_margin=Cm(0.9); section.bottom_margin=Cm(0.8); section.left_margin=Cm(1.0); section.right_margin=Cm(1.0)
    document.styles["Normal"].font.name="Aptos"; document.styles["Normal"].font.size=Pt(8)
    document.core_properties.author="EvoCowork AutomationBench 5121"; document.core_properties.title="Contractor Renewal Control Brief"; document.core_properties.created=datetime(2026,4,15,9,0,0); document.core_properties.modified=datetime(2026,4,15,9,0,0)
    add_doc_title(document,"Contractor Renewal Control Brief","AutomationBench 5121 · 60-day window · PIP and rate controls · draft only")
    add_body(document,"Eight source contractor agreements resolve to three RENEWED_STANDARD, two PENDING_VP_APPROVAL, one EXPIRED_NOT_RENEWED and two FUTURE_NO_ACTION. The frozen date is 2026-04-15. Expired takes precedence; otherwise only agreements ending in 0 through 60 days are processed, and active PIP cases cannot be renewed directly.")
    add_doc_table(document,["Decision","Count","Employees","Authority","Execution"],[
        ("RENEWED_STANDARD",3,"Ana Petrova; Darius Kim; Hugo Santos","Rules 1 and 3","3 notices / 3 updates / 0 executed"),
        ("PENDING_VP_APPROVAL",2,"Carmen Ruiz; Gloria Chen","Rule 2 active PIP","2 VP requests / 2 updates / 0 executed"),
        ("EXPIRED_NOT_RENEWED",1,"Felix Braun","Rule 5 expired","0 email / 1 update / 0 executed"),
        ("FUTURE_NO_ACTION",2,"Brian Osei; Eva Lindström","Rule 4 beyond 60 days","0 action"),
    ],[5.0,1.8,7.6,7.2,7.0])
    add_doc_table(document,["Rate control","Source fact","Applied decision","External state"],[
        ("User requested 5%","Original query","Escalate one consolidated request to Procurement","DRAFT_NOT_SENT"),
        ("HR Ops authority","Policy rule 6","No rate change for any contractor","0 applied"),
        ("Renewal notices","Policy rule 6","Existing terms only; never mention 5%","3 drafts / 0 sent"),
        ("Rate data","No rate field in tracker","No amount invented","0 invented"),
    ],[6.0,7.2,11.0,5.0])
    add_body(document,"The Procurement escalation covers only Ana Petrova, Darius Kim and Hugo Santos, asks Procurement to coordinate required Finance approval, and records that no rate was applied. Carmen and Gloria wait for VP approval; Felix is expired; Brian and Eva are outside the 60-day window. None belongs in the rate request.")
    document.add_page_break(); add_doc_title(document,"Complete eight-contractor decision register","Every source row retains end date, window, PIP, route, status and rate boundary")
    rows=[]
    for item in decisions:
        source=item.source; rows.append((source.source_row,source.contractor,source.department,source.agreement_end.isoformat(),item.days_to_end,item.on_pip,item.decision,item.proposed_status,item.to_email or "—",item.cc_email or "—",display_rate(item)))
    add_doc_table(document,["Row","Contractor","Dept.","End","Days","PIP","Decision","Proposed status","To","Cc","Rate treatment"],rows,[1.1,3.1,2.4,2.2,1.2,1.2,4.0,4.0,5.3,5.3,7.3],font_size=5.7)
    add_body(document,"Boundary cases are explicit: Ana has 30 days and is renewed; Gloria has 56 days but active PIP and routes only to vp-eng@company.example.com; Felix is -5 days and is not renewed; Brian is 76 days and Eva is 122 days, so both stay Active with zero draft actions.")
    document.add_page_break(); add_doc_title(document,"Controlled communication and tracker queue","Six source-backed email drafts · six tracker drafts · zero external execution")
    actions=[item for item in decisions if item.contractor_notice=="YES"]
    communication_rows=[]
    for index,item in enumerate(actions,1):
        source=item.source; communication_rows.append((index,"STANDARD_RENEWAL_NOTICE" if item.decision=="RENEWED_STANDARD" else "PIP_VP_APPROVAL_REQUEST",source.contractor,item.to_email,item.cc_email or "—",source.agreement_end.isoformat(),item.decision,"DRAFT_NOT_SENT"))
    communication_rows.append((6,"PROCUREMENT_RATE_ESCALATION","Ana; Darius; Hugo","procurement@company.example.com","—","—","5% REQUEST / NO RATE APPLIED","DRAFT_NOT_SENT"))
    add_doc_table(document,["Order","Type","Scope","To","Cc","End","Decision / request","State"],communication_rows,[1.1,5.2,3.1,5.6,5.2,2.3,5.2,3.6],font_size=5.9)
    add_doc_table(document,["Control","Draft / proposed","Executed","State"],[("Emails",6,0,"DRAFT_NOT_SENT"),("Tracker updates",6,0,"DRAFT_NOT_APPLIED"),("Rate changes",0,0,"NO_RATE_APPLIED"),("Slack",0,0,"NOT_AUTHORIZED_CHANNEL"),("Source workbook",0,0,"UNCHANGED"),("Real side effects",0,0,"OFFLINE_ONLY")],[7.0,5.0,4.0,10.0])
    add_body(document,"Standard notices go to each contractor with the hiring manager in Cc and confirm existing terms only. PIP requests go only to the department VP, with no Cc and no contractor contact. No Bcc, attachment or invented rate is allowed. #hr-general remains at zero drafts and zero posts because the frozen policy does not authorize Slack publication.")
    document.save(path); normalize_ooxml(path)


def add_textbox(slide,x,y,w,h,text,*,font_size=16,bold=False,color=PPT_NAVY,align=PP_ALIGN.LEFT,fill=None,line=None,radius=False):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); shape.fill.solid(); shape.fill.fore_color.rgb=fill or PPT_WHITE; shape.line.color.rgb=line or (fill or PPT_WHITE)
    frame=shape.text_frame; frame.clear(); frame.vertical_anchor=MSO_ANCHOR.MIDDLE; frame.margin_left=Inches(0.07); frame.margin_right=Inches(0.07); frame.margin_top=Inches(0.03); frame.margin_bottom=Inches(0.03)
    paragraph=frame.paragraphs[0]; paragraph.alignment=align; run=paragraph.add_run(); run.text=text; run.font.name="Aptos"; run.font.size=PptPt(font_size); run.font.bold=bold; run.font.color.rgb=color
    return shape


def add_slide_title(slide,title,subtitle):
    add_textbox(slide,0.55,0.25,12.25,0.55,title,font_size=24,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER)
    add_textbox(slide,0.75,0.78,11.85,0.36,subtitle,font_size=10.5,color=PptRGBColor(89,89,89),align=PP_ALIGN.CENTER)


def style_ppt_table(table,body_font=8.5):
    for row_index,row in enumerate(table.rows):
        for cell in row.cells:
            cell.fill.solid(); cell.fill.fore_color.rgb=PPT_NAVY if row_index==0 else (PPT_LIGHT_BLUE if row_index%2 else PPT_WHITE); cell.margin_left=Inches(0.04); cell.margin_right=Inches(0.04); cell.margin_top=Inches(0.02); cell.margin_bottom=Inches(0.02); cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment=PP_ALIGN.LEFT
                for run in paragraph.runs: run.font.name="Aptos"; run.font.size=PptPt(9 if row_index==0 else body_font); run.font.bold=row_index==0; run.font.color.rgb=PPT_WHITE if row_index==0 else PPT_NAVY


def write_presentation(path: Path, decisions: list[Decision]) -> None:
    presentation=Presentation(); presentation.slide_width=Inches(13.333); presentation.slide_height=Inches(7.5); blank=presentation.slide_layouts[6]
    presentation.core_properties.author="EvoCowork AutomationBench 5121"; presentation.core_properties.title="Contractor Renewal Review"; presentation.core_properties.created=datetime(2026,4,15,9,0,0); presentation.core_properties.modified=datetime(2026,4,15,9,0,0)
    slide=presentation.slides.add_slide(blank); add_slide_title(slide,"Contractor renewal control review","8 reviewed · 3 standard renewals · 2 PIP approvals · 1 expired · 2 future no action")
    cards=[("8","source contractors",PPT_LIGHT_BLUE),("3","standard renewals",PPT_GREEN),("2","PIP approvals",PPT_AMBER),("1","expired",PPT_RED),("2","future no action",PPT_LIGHT_BLUE)]
    for index,(value,label,fill) in enumerate(cards):
        left=0.55+index*2.5; add_textbox(slide,left,1.25,2.15,0.72,value,font_size=25,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=fill,line=PPT_BLUE,radius=True); add_textbox(slide,left,1.92,2.15,0.34,label,font_size=9,color=PPT_NAVY,align=PP_ALIGN.CENTER)
    table=slide.shapes.add_table(5,5,Inches(0.55),Inches(2.5),Inches(12.2),Inches(2.75)).table
    rows=[("Decision","Count","Contractors","Authority","Execution"),("RENEWED_STANDARD",3,"Ana; Darius; Hugo","rules 1 + 3","3 notices / 3 updates"),("PENDING_VP_APPROVAL",2,"Carmen; Gloria","rule 2 active PIP","2 VP requests / 2 updates"),("EXPIRED_NOT_RENEWED",1,"Felix","rule 5 expired","0 email / 1 update"),("FUTURE_NO_ACTION",2,"Brian; Eva","rule 4 beyond 60 days","0 action")]
    for r,values in enumerate(rows):
        for c,value in enumerate(values): table.cell(r,c).text=str(value)
    for index,width in enumerate([2.85,0.8,3.1,2.55,2.9]): table.columns[index].width=Inches(width)
    style_ppt_table(table,8.4)
    add_textbox(slide,0.55,5.62,12.2,0.62,"5% is a request, not an applied rate: one Procurement escalation · Finance approval required · no rate invented",font_size=11.4,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_AMBER,line=PPT_AMBER,radius=True)
    add_textbox(slide,0.55,6.42,12.2,0.42,"Precedence: expired -> beyond 60 days -> active PIP -> standard renewal · boundaries are inclusive 0–60 days",font_size=10.2,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_LIGHT_BLUE,line=PPT_LIGHT_BLUE,radius=True)

    slide=presentation.slides.add_slide(blank); add_slide_title(slide,"Complete eight-contractor decision register","Frozen date 2026-04-15 · end date · days · PIP · route · rate boundary")
    table=slide.shapes.add_table(9,10,Inches(0.25),Inches(1.2),Inches(12.83),Inches(5.6)).table
    headers=("Row","Contractor","Dept.","End","Days","PIP","Decision","Status","To","Rate")
    for c,value in enumerate(headers): table.cell(0,c).text=value
    for r,item in enumerate(decisions,1):
        source=item.source; values=(source.source_row,source.contractor,source.department,source.agreement_end.isoformat(),item.days_to_end,item.on_pip,item.decision,item.proposed_status,item.to_email or "—",display_rate(item))
        for c,value in enumerate(values): table.cell(r,c).text=str(value)
    for index,width in enumerate([0.45,1.3,0.75,1.0,0.5,0.48,1.62,1.42,2.3,3.01]): table.columns[index].width=Inches(width)
    style_ppt_table(table,6.5)
    add_textbox(slide,0.45,6.88,12.43,0.34,"Ana: 30-day standard · Gloria: 56-day PIP to VP · Felix: -5-day expired · Brian/Eva: 76/122-day no action",font_size=9.4,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_LIGHT_BLUE,line=PPT_LIGHT_BLUE,radius=True)

    slide=presentation.slides.add_slide(blank); add_slide_title(slide,"Controlled action queue","6 source-backed emails · 6 tracker drafts · no rate change · no Slack · zero execution")
    actions=[item for item in decisions if item.contractor_notice=="YES"]
    table=slide.shapes.add_table(7,8,Inches(0.25),Inches(1.25),Inches(12.83),Inches(3.65)).table
    headers=("Order","Type","Scope","To","Cc","End","Decision / request","State")
    for c,value in enumerate(headers): table.cell(0,c).text=value
    for r,item in enumerate(actions,1):
        source=item.source; values=(r,"STANDARD_NOTICE" if item.decision=="RENEWED_STANDARD" else "PIP_VP_REQUEST",source.contractor,item.to_email,item.cc_email or "—",source.agreement_end.isoformat(),item.decision,"DRAFT_ONLY")
        for c,value in enumerate(values): table.cell(r,c).text=str(value)
    procurement=(6,"PROCUREMENT_ESCALATION","Ana; Darius; Hugo","procurement@company.example.com","—","—","5% request / no rate applied","DRAFT_ONLY")
    for c,value in enumerate(procurement): table.cell(6,c).text=str(value)
    for index,width in enumerate([0.5,1.55,1.45,2.7,2.35,0.95,2.05,1.28]): table.columns[index].width=Inches(width)
    style_ppt_table(table,6.6)
    add_textbox(slide,0.55,5.15,3.8,0.72,"Email\n6 drafts · 0 sent",font_size=13,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_GREEN,line=PPT_BLUE,radius=True)
    add_textbox(slide,4.77,5.15,3.8,0.72,"Tracker\n6 drafts · 0 applied",font_size=13,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_AMBER,line=PPT_BLUE,radius=True)
    add_textbox(slide,8.99,5.15,3.8,0.72,"Rate / Slack\n0 applied · 0 posts",font_size=13,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_GREEN,line=PPT_BLUE,radius=True)
    add_textbox(slide,0.55,6.08,12.2,0.5,"Standard: contractor To + manager Cc + existing terms only · PIP: department VP To only · Bcc=0 · attachments=0",font_size=9.8,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_LIGHT_BLUE,line=PPT_LIGHT_BLUE,radius=True)
    add_textbox(slide,0.55,6.76,12.2,0.4,"Source workbook changes=0 · real side effects=0 · #hr-general drafts=0 and posts=0",font_size=9.8,bold=True,color=PPT_NAVY,align=PP_ALIGN.CENTER,fill=PPT_AMBER,line=PPT_AMBER,radius=True)
    presentation.save(path); normalize_ooxml(path)


def email_body(item: Decision) -> str:
    source=item.source
    if item.decision == "RENEWED_STANDARD":
        return "\n".join([
            f"Hello {source.contractor},", "",
            f"Your contractor agreement ending {source.agreement_end.isoformat()} is within the 60-day renewal window.",
            "This draft confirms continuation at existing terms only. No rate adjustment has been applied or promised.",
            f"Your hiring manager, {source.manager}, is copied for confirmation.",
            "Proposed tracker status: Renewed.", "",
            "DRAFT_NOT_SENT — this offline benchmark artifact has not been sent, and the source tracker has not been changed.",
            "- HR Operations",
        ])
    if item.decision != "PENDING_VP_APPROVAL":
        raise ValueError(f"{source.contractor} 不应生成候选人通知")
    return "\n".join([
        "Hello Department VP,", "",
        f"Approval is required before renewing {source.contractor}'s contractor agreement ending {source.agreement_end.isoformat()}.",
        f"The tracker records an active PIP: {source.notes}.",
        "The contractor has not been contacted, and no rate action has been taken.",
        "Proposed tracker status: Pending VP Approval.", "",
        "DRAFT_NOT_SENT — this offline benchmark artifact has not been sent, and the source tracker has not been changed.",
        "- HR Operations",
    ])


def procurement_body() -> str:
    return "\n".join([
        "Hello Procurement,", "",
        "The original renewal request asked HR Ops to apply a standard 5% annual rate increase.",
        "Policy prohibits HR Ops from changing contractor rates and requires Procurement and Finance approval.",
        "Please review the 5% request for these standard-renewal candidates only:",
        "- Ana Petrova — agreement end 2026-05-15",
        "- Darius Kim — agreement end 2026-05-01",
        "- Hugo Santos — agreement end 2026-05-20",
        "", "No rate amount exists in the source tracker, no rate has been invented or applied, and renewal notices continue existing terms only.",
        "Please coordinate the required Finance approval before any later rate action.", "",
        "DRAFT_NOT_SENT — this offline benchmark artifact has not been sent, and the source tracker has not been changed.",
        "- HR Operations",
    ])


def write_mbox(path: Path, decisions: list[Decision]) -> None:
    chunks=[]; order=0
    for item in decisions:
        if item.contractor_notice != "YES": continue
        order += 1; source=item.source; message=EmailMessage(policy=SMTP); message["From"]="hr-ops@company.example.com"; message["To"]=item.to_email
        if item.cc_email: message["Cc"]=item.cc_email
        message["Subject"] = f"Contractor Renewal at Existing Terms - {source.contractor}" if item.decision=="RENEWED_STANDARD" else f"VP Approval Required for Contractor Renewal - {source.contractor}"
        message["Date"]="Wed, 15 Apr 2026 09:00:00 +0000"; message["Message-ID"]=f"<contractor-5121-{source.source_row}@evocowork.local>"; message["X-EvoCowork-Status"]="DRAFT_NOT_SENT"; message["X-EvoCowork-Source-Row"]=str(source.source_row); message["X-EvoCowork-Decision"]=item.decision; message["X-EvoCowork-Proposed-Status"]=item.proposed_status; message["X-EvoCowork-Rate-Change-Applied"]="NO"; message["X-EvoCowork-Order"]=str(order)
        message.set_content(email_body(item),charset="utf-8",cte="quoted-printable"); stream=BytesIO(); BytesGenerator(stream,policy=SMTP.clone(linesep="\n",max_line_length=78)).flatten(message,unixfrom=False)
        chunks.append(f"From MAILER-DAEMON Wed Apr 15 09:00:{order:02d} 2026\n".encode("ascii")+stream.getvalue().replace(b"\nFrom ",b"\n>From ")+b"\n")
    order += 1
    message=EmailMessage(policy=SMTP); message["From"]="hr-ops@company.example.com"; message["To"]="procurement@company.example.com"; message["Subject"]="Rate Change Approval Request for Three Contractor Renewals"; message["Date"]="Wed, 15 Apr 2026 09:00:00 +0000"; message["Message-ID"]="<contractor-5121-procurement@evocowork.local>"; message["X-EvoCowork-Status"]="DRAFT_NOT_SENT"; message["X-EvoCowork-Decision"]="RATE_CHANGE_APPROVAL_REQUEST"; message["X-EvoCowork-Rate-Change-Applied"]="NO"; message["X-EvoCowork-Order"]=str(order); message["X-EvoCowork-Contractor-Scope"]="Ana Petrova; Darius Kim; Hugo Santos"
    message.set_content(procurement_body(),charset="utf-8",cte="quoted-printable"); stream=BytesIO(); BytesGenerator(stream,policy=SMTP.clone(linesep="\n",max_line_length=78)).flatten(message,unixfrom=False)
    chunks.append(f"From MAILER-DAEMON Wed Apr 15 09:00:{order:02d} 2026\n".encode("ascii")+stream.getvalue().replace(b"\nFrom ",b"\n>From ")+b"\n")
    if order != 6: raise RuntimeError("续签邮件必须恰好五封个案草稿加一封 Procurement 升级")
    path.write_bytes(b"".join(chunks))


def write_updates(path: Path, decisions: list[Decision]) -> None:
    drafts=[]
    for item in decisions:
        if item.tracker_update != "YES": continue
        source=item.source; drafts.append({"source_row":source.source_row,"contractor":source.contractor,"agreement_end":source.agreement_end.isoformat(),"days_to_end":item.days_to_end,"on_pip":item.on_pip,"decision":item.decision,"source_status":source.source_status,"proposed_status":item.proposed_status,"rate_change_applied":False,"rate_treatment":item.rate_treatment,"state":"DRAFT_NOT_APPLIED","applied":False,"source_trace":[f"Contractor_Agreement_Tracker.xlsx:Active Agreements row {source.source_row}","Contractor_Renewal_Policy.mbox:msg_contractor_policy_5121","Workflow_Context.json:meta.current_time"]})
    payload={"task_id":"automationbench-5121","status":"CONTROLLED_DRAFT_ONLY","frozen_time":FIXED_ISO_TIME,"window_method":"calendar days from 2026-04-15; inclusive 0 through 60 days","precedence":["EXPIRED_NOT_RENEWED","FUTURE_NO_ACTION","PENDING_VP_APPROVAL","RENEWED_STANDARD"],"draft_count":6,"applied_count":0,"source_modified":False,"rate_changes_applied":0,"drafts":drafts}
    path.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")


def write_slack(path: Path) -> None:
    payload={"task_id":"automationbench-5121","source_channel":{"id":"C_HR_5121","name":"hr-general","is_private":False},"decision":"NOT_AUTHORIZED_CHANNEL","reason":"The frozen contractor renewal policy authorizes email and tracker actions but does not authorize Slack publication.","draft_count":0,"posted_count":0,"posted":False,"contractor_names_in_slack":[],"pip_details_in_slack":[],"rate_details_in_slack":[],"messages_created":[],"source_trace":["Contractor_Renewal_Policy.mbox:msg_contractor_policy_5121","Slack_Channels.json:C_HR_5121"]}
    path.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")


def write_execution(path: Path) -> None:
    payload={"task_id":"automationbench-5121","status":"CONTROLLED_DRAFT_ONLY","email":{"standard_renewal_drafts":3,"pip_vp_approval_drafts":2,"procurement_rate_escalation_drafts":1,"total_drafts":6,"sent_count":0,"state":"DRAFT_NOT_SENT"},"tracker_updates":{"standard_renewal_drafts":3,"pip_pending_approval_drafts":2,"expired_not_renewed_drafts":1,"total_drafts":6,"applied_count":0,"source_modified":False,"state":"DRAFT_NOT_APPLIED"},"rate_control":{"requested_in_query":"5% annual increase","approval_route":"Procurement and Finance","source_rate_fields":0,"invented_rate_amounts":0,"rate_changes_applied":0,"renewal_terms":"EXISTING_TERMS_ONLY","state":"REQUEST_ESCALATED_NOT_APPLIED"},"no_action":{"future_contractors":["Brian Osei","Eva Lindström"],"email_draft_count":0,"tracker_update_draft_count":0},"slack":{"channel":"C_HR_5121 / hr-general","draft_count":0,"posted_count":0,"state":"NOT_AUTHORIZED_CHANNEL"},"source_tracker_modification_count":0,"real_side_effect_count":0,"source_trace":["Contractor_Agreement_Tracker.xlsx:Active Agreements rows 2-9","Contractor_Renewal_Policy.mbox:msg_contractor_policy_5121","Workflow_Context.json:meta.current_time","Slack_Channels.json:C_HR_5121"]}
    path.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")


def main() -> None:
    parser=argparse.ArgumentParser(); parser.add_argument("--input-dir",type=Path,required=True); parser.add_argument("--output-dir",type=Path,required=True); args=parser.parse_args()
    records,policy_message,channels,frozen_time=read_inputs(args.input_dir); decisions=derive_decisions(records)
    args.output_dir.mkdir(parents=True,exist_ok=True)
    for item in args.output_dir.iterdir():
        if item.is_file(): item.unlink()
        elif item.is_dir(): shutil.rmtree(item)
    write_workbook(args.output_dir/XLSX_NAME,records,policy_message,channels,frozen_time); write_brief(args.output_dir/DOCX_NAME,decisions); write_presentation(args.output_dir/PPTX_NAME,decisions); write_mbox(args.output_dir/MBOX_NAME,decisions); write_updates(args.output_dir/UPDATES_NAME,decisions); write_slack(args.output_dir/SLACK_NAME); write_execution(args.output_dir/EXECUTION_NAME)
    expected={XLSX_NAME,DOCX_NAME,PPTX_NAME,MBOX_NAME,UPDATES_NAME,SLACK_NAME,EXECUTION_NAME}
    if {item.name for item in args.output_dir.iterdir() if item.is_file()} != expected: raise RuntimeError("交付文件集合与冻结任务不一致")


if __name__ == "__main__":
    main()
