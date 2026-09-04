#!/usr/bin/env python3
"""AutomationBench 5121 的任务专属 v2 分层负例。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from typing import Callable

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
INPUTS = ROOT / "environment" / "input_files"
OUTPUTS = (
    "Contractor_Renewal_Audit.xlsx", "Contractor_Renewal_Control_Brief.docx",
    "Contractor_Renewal_Review.pptx", "Contractor_Renewal_Notification_Drafts.mbox",
    "Contractor_Tracker_Update_Drafts.json", "Contractor_Slack_Control.json",
    "Contractor_Execution_Control.json",
)
Mutator = Callable[[Path], None]


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"无法载入{path}")
    module = importlib.util.module_from_spec(spec); sys.modules[spec.name] = module; spec.loader.exec_module(module)
    return module


def _xlsx(root: Path, sheet: str, cell: str, value) -> None:
    path = root / OUTPUTS[0]; book = load_workbook(path); book[sheet][cell] = value; book.save(path)


def _json(root: Path, name: str, mutator: Callable[[dict], None]) -> None:
    path = root / name; payload = json.loads(path.read_text(encoding="utf-8")); mutator(payload)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _rewrite_mbox(root: Path, mutator: Callable[[list], None]) -> None:
    path = root / OUTPUTS[3]; box = mailbox.mbox(path, create=False); messages = list(box); box.close(); mutator(messages)
    with path.open("wb") as handle:
        for index, message in enumerate(messages, 1):
            handle.write(f"From MAILER-DAEMON Wed Apr 15 09:00:{index:02d} 2026\n".encode())
            BytesGenerator(handle, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message); handle.write(b"\n")


def _replace_word(root: Path, old: str, new: str) -> None:
    path = root / OUTPUTS[1]; document = Document(path); changed = False
    for paragraph in document.paragraphs:
        if old in paragraph.text:
            paragraph.text = paragraph.text.replace(old, new, 1); changed = True; break
    if not changed:
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if old in cell.text:
                        cell.text = cell.text.replace(old, new, 1); changed = True; break
                if changed: break
            if changed: break
    if not changed: raise ValueError(f"Word中找不到{old!r}")
    document.save(path)


def _replace_ppt(root: Path, old: str, new: str) -> None:
    path = root / OUTPUTS[2]; deck = Presentation(path); changed = False
    for slide in deck.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and old in shape.text:
                shape.text = shape.text.replace(old, new, 1); changed = True
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        if old in cell.text:
                            cell.text = cell.text.replace(old, new, 1); changed = True; break
                    if changed: break
            if changed: break
        if changed: break
    if not changed: raise ValueError(f"PPT中找不到{old!r}")
    deck.save(path)


def _input_only(root: Path) -> None:
    for path in list(root.iterdir()):
        if path.is_file(): path.unlink()
    for source in INPUTS.iterdir():
        if source.is_file(): shutil.copy2(source, root / source.name)


def _promise_rate(messages: list) -> None:
    body = messages[0].get_payload(decode=True).decode(messages[0].get_content_charset() or "utf-8")
    changed = body.replace("No rate adjustment has been applied or promised.", "A 5% annual rate increase will be applied.").replace("—", "-")
    messages[0].set_payload(changed)


def _cases() -> list[tuple[str, Mutator, str]]:
    cases: list[tuple[str, Mutator, str]] = []
    for name in OUTPUTS:
        cases.append((f"missing-{name}", lambda root, item=name: (root / item).unlink(), "fatal_zero"))
        cases.append((f"corrupt-{name}", lambda root, item=name: (root / item).write_bytes(b"not-a-valid-artifact"), "fatal_zero"))
    cases.extend([
        ("mbox-empty", lambda root: (root / OUTPUTS[3]).write_bytes(b""), "fatal_zero"),
        ("unmodified-input-only", _input_only, "fatal_zero"),
        ("source-ana-department-local", lambda root: _xlsx(root, "Source_Contractors", "F2", "Product"), "partial"),
        ("ana-asof-formula-static", lambda root: _xlsx(root, "Renewal_Decisions", "K2", "2026-04-15"), "partial"),
        ("felix-decision-wrong", lambda root: _xlsx(root, "Renewal_Decisions", "Q7", "RENEWED_STANDARD"), "partial"),
        ("mail-ana-subject-local", lambda root: _rewrite_mbox(root, lambda messages: messages[0].replace_header("Subject", "Renewal Notice - Ana Petrova")), "partial"),
        ("mail-carmen-contractor-recipient", lambda root: _rewrite_mbox(root, lambda messages: messages[1].replace_header("To", "carmen.ruiz@freelance.example.com")), "partial"),
        ("mail-ana-promises-five-percent", lambda root: _rewrite_mbox(root, _promise_rate), "partial"),
        ("mail-procurement-missing-finance", lambda root: _rewrite_mbox(root, lambda messages: messages[5].set_payload(messages[5].get_payload().replace("Procurement and Finance approval", "Procurement approval"))), "partial"),
        ("update-hugo-status-local", lambda root: _json(root, OUTPUTS[4], lambda payload: payload["drafts"][5].update({"proposed_status":"Active"})), "partial"),
        ("update-brian-added", lambda root: _json(root, OUTPUTS[4], lambda payload: payload["drafts"].append({"contractor":"Brian Osei"})), "partial"),
        ("slack-pip-rate-leak", lambda root: _json(root, OUTPUTS[5], lambda payload: payload.update({"draft_count":1,"contractor_names_in_slack":["Carmen Ruiz"],"pip_details_in_slack":["On PIP"],"rate_details_in_slack":["5%"],"messages_created":[{"text":"Carmen Ruiz is on PIP; 5% rate"}]})), "partial"),
        ("execution-email-sent", lambda root: _json(root, OUTPUTS[6], lambda payload: payload["email"].update({"sent_count":1,"state":"SENT"})), "partial"),
        ("execution-rate-applied", lambda root: _json(root, OUTPUTS[6], lambda payload: payload["rate_control"].update({"rate_changes_applied":1,"state":"APPLIED"})), "partial"),
        ("word-felix-boundary-local", lambda root: _replace_word(root, "Felix is -5 days", "Felix is 5 days"), "partial"),
        ("ppt-gloria-boundary-local", lambda root: _replace_ppt(root, "Gloria: 56-day PIP to VP", "Gloria: standard renewal"), "partial"),
        ("summary-standard-count-local", lambda root: _xlsx(root, "Control_Summary", "B3", 2), "partial"),
        ("extra-file", lambda root: (root / "unexpected.txt").write_text("unexpected", encoding="utf-8"), "partial"),
    ])
    return cases


def main() -> None:
    parser = argparse.ArgumentParser(); parser.add_argument("--report", type=Path); args = parser.parse_args()
    judge = _load("contractor_5121_v2_judge", ROOT / "tests" / "judge.py")
    judge.V1.render = lambda path: ({".xlsx":10,".docx":3,".pptx":3}[path.suffix.casefold()], "render-ok")
    if judge.evaluate(GOLDEN).get("reward") != 1.0: raise AssertionError("Golden v2必须为1.0")
    results = []
    for name, mutator, layer in _cases():
        with tempfile.TemporaryDirectory(prefix=f"contractor-5121-v2-{name}-") as temporary:
            output = Path(temporary) / "output"; shutil.copytree(GOLDEN, output); mutator(output); result = judge.evaluate(output); reward = float(result.get("reward", 0))
            if layer == "fatal_zero" and reward != 0: raise AssertionError(f"{name}应硬归零，实际{reward}")
            if layer == "partial" and not (0 < reward < 1): raise AssertionError(f"{name}应保留连续部分分，实际{reward}")
            results.append({"case":name,"layer":layer,"reward":reward,"passed":bool(result.get("passed")),"penalties":result.get("penalties",[])})
    partial = [item["reward"] for item in results if item["layer"] == "partial"]
    report = {"total":len(results),"fatal_zero":sum(item["layer"]=="fatal_zero" for item in results),"partial_continuous":len(partial),"partial_reward_range":[min(partial),max(partial)],"cases":results}
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True); args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
