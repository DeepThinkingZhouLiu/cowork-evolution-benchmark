#!/usr/bin/env python3
"""AutomationBench 5122 的任务专属 v2 分层负例。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from typing import Callable

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]
GOLDEN=ROOT/"solution"/"Golden"
INPUTS=ROOT/"environment"/"input_files"
OUTPUTS=("Survey_Distribution_Audit.xlsx","Survey_Distribution_Brief.docx","Survey_Distribution_Review.pptx","Survey_Invitation_Drafts.mbox","All_Hands_Announcement_Draft.json","Survey_Delivery_Gap_Register.json","Survey_Execution_Control.json")
Mutator=Callable[[Path],None]


def _load(name: str,path: Path):
    spec=importlib.util.spec_from_file_location(name,path)
    if spec is None or spec.loader is None:raise RuntimeError(f"无法载入{path}")
    module=importlib.util.module_from_spec(spec);sys.modules[spec.name]=module;spec.loader.exec_module(module);return module


def _xlsx(root: Path,sheet: str,cell: str,value) -> None:
    path=root/OUTPUTS[0];book=load_workbook(path);book[sheet][cell]=value;book.save(path)


def _json(root: Path,name: str,mutator: Callable[[dict],None]) -> None:
    path=root/name;payload=json.loads(path.read_text(encoding="utf-8"));mutator(payload);path.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")


def _rewrite_mbox(root: Path,mutator: Callable[[list],None]) -> None:
    path=root/OUTPUTS[3];box=mailbox.mbox(path,create=False);messages=list(box);box.close();mutator(messages)
    with path.open("wb") as handle:
        for index,message in enumerate(messages,1):
            handle.write(f"From MAILER-DAEMON Wed Apr 15 09:00:{index:02d} 2026\n".encode());BytesGenerator(handle,policy=SMTP,mangle_from_=True,maxheaderlen=0).flatten(message);handle.write(b"\n")


def _replace_word(root: Path,old: str,new: str) -> None:
    path=root/OUTPUTS[1];document=Document(path);changed=False
    for paragraph in document.paragraphs:
        if old in paragraph.text:paragraph.text=paragraph.text.replace(old,new,1);changed=True;break
    if not changed:
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if old in cell.text:cell.text=cell.text.replace(old,new,1);changed=True;break
                if changed:break
            if changed:break
    if not changed:raise ValueError(f"Word中找不到{old!r}")
    document.save(path)


def _replace_ppt(root: Path,old: str,new: str) -> None:
    path=root/OUTPUTS[2];deck=Presentation(path);changed=False
    for slide in deck.slides:
        for shape in slide.shapes:
            if getattr(shape,"has_text_frame",False) and old in shape.text:shape.text=shape.text.replace(old,new,1);changed=True
            if getattr(shape,"has_table",False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        if old in cell.text:cell.text=cell.text.replace(old,new,1);changed=True;break
                    if changed:break
            if changed:break
        if changed:break
    if not changed:raise ValueError(f"PPT中找不到{old!r}")
    deck.save(path)


def _input_only(root: Path) -> None:
    for path in list(root.iterdir()):
        if path.is_file():path.unlink()
    for source in INPUTS.iterdir():
        if source.is_file():shutil.copy2(source,root/source.name)


def _remove_clara_invite(root: Path) -> None:
    path=root/OUTPUTS[0];book=load_workbook(path);sheet=book["Invitation_Drafts"];sheet.delete_rows(3,1);book.save(path)


def _add_clara_excluded(root: Path) -> None:
    path=root/OUTPUTS[0];book=load_workbook(path);sheet=book["Excluded_Employees"];sheet.append(["Clara Obi","Legal Operations","Full-time","Active","EXCLUDE_LEGAL_OVERRIDE"]);book.save(path)


def _cases() -> list[tuple[str,Mutator,str]]:
    cases=[]
    for name in OUTPUTS:
        cases.append((f"missing-{name}",lambda root,item=name:(root/item).unlink(),"fatal_zero"))
        cases.append((f"corrupt-{name}",lambda root,item=name:(root/item).write_bytes(b"not-a-valid-artifact"),"fatal_zero"))
    cases.extend([
        ("mbox-empty",lambda root:(root/OUTPUTS[3]).write_bytes(b""),"fatal_zero"),
        ("unmodified-input-only",_input_only,"fatal_zero"),
        ("source-clara-department-local",lambda root:_xlsx(root,"Source_Roster","D4","Legal"),"partial"),
        ("formula-clara-fuzzy-legal",lambda root:_xlsx(root,"Eligibility_Decisions","J4",'=IF(ISNUMBER(SEARCH("Legal",D4)),"YES","NO")'),"partial"),
        ("invite-clara-row-omitted",_remove_clara_invite,"partial"),
        ("excluded-clara-added",_add_clara_excluded,"partial"),
        ("mail-anna-subject-local",lambda root:_rewrite_mbox(root,lambda messages:messages[0].replace_header("Subject","Annual survey invitation - Anna Chen")),"partial"),
        ("mail-boris-wrong-recipient",lambda root:_rewrite_mbox(root,lambda messages:messages[0].replace_header("To","boris.petrov@company.example.com")),"partial"),
        ("mail-anna-wrong-link",lambda root:_rewrite_mbox(root,lambda messages:messages[0].set_payload(messages[0].get_payload().replace("survey.company.example.com/2026","wrong.example.com/2026"))),"partial"),
        ("announcement-wrong-channel",lambda root:_json(root,OUTPUTS[4],lambda payload:payload.update({"channel_id":"C_HR_5122"})),"partial"),
        ("announcement-posted",lambda root:_json(root,OUTPUTS[4],lambda payload:payload.update({"state":"POSTED","posted_count":1})),"partial"),
        ("gap-elena-owner-local",lambda root:_json(root,OUTPUTS[5],lambda payload:payload["gaps"][0].update({"owner":"Unassigned"})),"partial"),
        ("gap-elena-omitted",lambda root:_json(root,OUTPUTS[5],lambda payload:payload.update({"gap_count":0,"gaps":[]})),"partial"),
        ("execution-eligible-count",lambda root:_json(root,OUTPUTS[6],lambda payload:payload.update({"eligible_employees":5})),"partial"),
        ("execution-emails-sent",lambda root:_json(root,OUTPUTS[6],lambda payload:payload.update({"emails_sent":5,"real_side_effects":5})),"partial"),
        ("word-fuzzy-legal-exclusion",lambda root:_replace_word(root,"Clara Obi in Legal Operations remains included.","Clara Obi was excluded because Legal Operations contains Legal."),"partial"),
        ("ppt-elena-gap-omitted",lambda root:_replace_ppt(root,"Elena Rossi","Unlisted employee"),"partial"),
        ("extra-file",lambda root:(root/"unexpected.txt").write_text("unexpected",encoding="utf-8"),"partial"),
    ])
    return cases


def main() -> None:
    parser=argparse.ArgumentParser();parser.add_argument("--report",type=Path);args=parser.parse_args();judge=_load("survey_5122_v2_judge",ROOT/"tests"/"judge.py");judge.V1.render=lambda path:({".xlsx":10,".docx":3,".pptx":3}[path.suffix.casefold()],"render-ok")
    if judge.evaluate(GOLDEN).get("reward")!=1.0:raise AssertionError("Golden v2必须为1.0")
    results=[]
    for name,mutator,layer in _cases():
        with tempfile.TemporaryDirectory(prefix=f"survey-5122-v2-{name}-") as temporary:
            output=Path(temporary)/"output";shutil.copytree(GOLDEN,output);mutator(output);result=judge.evaluate(output);reward=float(result.get("reward",0))
            if layer=="fatal_zero" and reward!=0:raise AssertionError(f"{name}应硬归零，实际{reward}")
            if layer=="partial" and not (0<reward<1):raise AssertionError(f"{name}应保留连续部分分，实际{reward}")
            results.append({"case":name,"layer":layer,"reward":reward,"passed":bool(result.get("passed")),"penalties":result.get("penalties",[])})
    partial=[item["reward"] for item in results if item["layer"]=="partial"]
    report={"total":len(results),"fatal_zero":sum(item["layer"]=="fatal_zero" for item in results),"partial_continuous":len(partial),"partial_reward_range":[min(partial),max(partial)],"cases":results}
    if args.report:args.report.parent.mkdir(parents=True,exist_ok=True);args.report.write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=="__main__":main()
