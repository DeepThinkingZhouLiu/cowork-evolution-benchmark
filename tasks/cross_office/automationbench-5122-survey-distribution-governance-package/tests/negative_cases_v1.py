#!/usr/bin/env python3
import json,mailbox,shutil,subprocess,sys,tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1];GOLDEN=ROOT/"solution/Golden";PYTHON=sys.executable
def score(case):
 result=case.parent/(case.name+".json");reward=case.parent/(case.name+".txt");subprocess.run([PYTHON,str(ROOT/"tests/judge.py"),"--artifact-dir",str(case),"--result",str(result),"--reward-file",str(reward)],check=True,capture_output=True,text=True,timeout=240);return float(reward.read_text())
def rewrite(path,mutation):
 box=mailbox.mbox(path,create=False);messages=list(box);box.close();mutation(messages)
 with path.open("wb") as handle:
  for index,message in enumerate(messages):handle.write(f"From x Wed Apr 15 09:00:{index:02d} 2026\n".encode());BytesGenerator(handle,policy=SMTP,mangle_from_=True,maxheaderlen=0).flatten(message);handle.write(b"\n")
def replace_body(message,text):
 for header in ("Content-Type","Content-Transfer-Encoding","MIME-Version"):
  if header in message:del message[header]
 message.set_payload(text)
def main():
 with tempfile.TemporaryDirectory(prefix="survey5122-neg-") as temp:
  root=Path(temp);names=["missing","wrong_count","legal_invited","clara_excluded","elena_omitted","wrong_link","wrong_channel","wrong_gap","emails_sent","formula_removed","word_false","ppt_omit"];cases=[]
  for name in names:case=root/name;shutil.copytree(GOLDEN,case);cases.append(case)
  (cases[0]/"Survey_Distribution_Review.pptx").unlink()
  payload=json.loads((cases[1]/"Survey_Execution_Control.json").read_text());payload["eligible_employees"]=5;(cases[1]/"Survey_Execution_Control.json").write_text(json.dumps(payload))
  rewrite(cases[2]/"Survey_Invitation_Drafts.mbox",lambda ms:ms[0].replace_header("To","boris.petrov@company.example.com"))
  rewrite(cases[3]/"Survey_Invitation_Drafts.mbox",lambda ms:ms.pop(1))
  payload=json.loads((cases[4]/"Survey_Delivery_Gap_Register.json").read_text());payload["gap_count"]=0;payload["gaps"]=[];(cases[4]/"Survey_Delivery_Gap_Register.json").write_text(json.dumps(payload))
  rewrite(cases[5]/"Survey_Invitation_Drafts.mbox",lambda ms:replace_body(ms[0],"Hello Anna Chen,\n\nPlease use https://wrong.example.com.\n\nDRAFT_NOT_SENT."))
  payload=json.loads((cases[6]/"All_Hands_Announcement_Draft.json").read_text());payload["channel_id"]="C_HR";(cases[6]/"All_Hands_Announcement_Draft.json").write_text(json.dumps(payload))
  payload=json.loads((cases[7]/"Survey_Delivery_Gap_Register.json").read_text());payload["gaps"][0]["employee"]="Unknown";(cases[7]/"Survey_Delivery_Gap_Register.json").write_text(json.dumps(payload))
  payload=json.loads((cases[8]/"Survey_Execution_Control.json").read_text());payload["emails_sent"]=5;(cases[8]/"Survey_Execution_Control.json").write_text(json.dumps(payload))
  workbook=load_workbook(cases[9]/"Survey_Distribution_Audit.xlsx");workbook["Eligibility_Decisions"]["M2"]="INVITATION_DRAFT";workbook.save(cases[9]/"Survey_Distribution_Audit.xlsx")
  document=Document(cases[10]/"Survey_Distribution_Brief.docx");document.add_paragraph("Clara Obi was excluded because Legal Operations contains Legal.");document.save(cases[10]/"Survey_Distribution_Brief.docx")
  presentation=Presentation(cases[11]/"Survey_Distribution_Review.pptx")
  for slide in presentation.slides:
   for shape in slide.shapes:
    if getattr(shape,"has_text_frame",False):shape.text=shape.text.replace("Elena Rossi","Unknown")
    if getattr(shape,"has_table",False):
     for row in shape.table.rows:
      for cell in row.cells:cell.text=cell.text.replace("Elena Rossi","Unknown")
  presentation.save(cases[11]/"Survey_Distribution_Review.pptx")
  rewards={case.name:score(case) for case in cases};print(json.dumps(rewards,indent=2));assert all(value==0 for value in rewards.values()),rewards
if __name__=="__main__":main()
