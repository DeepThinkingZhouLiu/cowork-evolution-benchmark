#!/usr/bin/env python3
"""逐项注入 AutomationBench 5123 专属错误，验证 Judge 不会误接受坏答案。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
import zipfile
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path
from typing import Callable

from docx import Document
from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pptx import Presentation


ROOT=Path(__file__).resolve().parents[1]; GOLDEN=ROOT/"solution"/"Golden"
XLSX="Job_Posting_Audit.xlsx"; DOCX="Job_Posting_Control_Brief.docx"; PPTX="Job_Posting_Review.pptx"; MBOX="Job_Posting_Email_Drafts.mbox"; ACTIONS="Job_Posting_Action_Drafts.json"; SLACK="Job_Posting_Slack_Control.json"; EXECUTION="Job_Posting_Execution_Control.json"
FILES=(XLSX,DOCX,PPTX,MBOX,ACTIONS,SLACK,EXECUTION); Mutator=Callable[[Path],None]; Case=tuple[str,Mutator,str]


def load_judge():
    spec=importlib.util.spec_from_file_location("job_posting_5123_judge",ROOT/"tests"/"judge.py")
    if spec is None or spec.loader is None: raise RuntimeError("无法载入本题 Judge")
    module=importlib.util.module_from_spec(spec); sys.modules[spec.name]=module; spec.loader.exec_module(module); module.render=lambda path:({".xlsx":11,".docx":3,".pptx":3}[path.suffix],"render-ok"); return module


def rewrite_zip(path: Path,transform: Callable[[str,bytes],bytes|None]) -> None:
    temporary=path.with_suffix(path.suffix+".negative")
    with zipfile.ZipFile(path,"r") as source,zipfile.ZipFile(temporary,"w",zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data=transform(info.filename,source.read(info.filename))
            if data is not None: target.writestr(info,data)
    temporary.replace(path)


def sheet_xml_path(archive: zipfile.ZipFile,sheet_name: str) -> str:
    workbook=etree.fromstring(archive.read("xl/workbook.xml")); relationships=etree.fromstring(archive.read("xl/_rels/workbook.xml.rels")); ns={"m":"http://schemas.openxmlformats.org/spreadsheetml/2006/main","r":"http://schemas.openxmlformats.org/officeDocument/2006/relationships","pr":"http://schemas.openxmlformats.org/package/2006/relationships"}
    relationship_id=workbook.xpath(f'//m:sheet[@name="{sheet_name}"]/@r:id',namespaces=ns)[0]; target=relationships.xpath(f'//pr:Relationship[@Id="{relationship_id}"]/@Target',namespaces=ns)[0]; return target.lstrip("/") if target.startswith("/") else "xl/"+target.lstrip("/")


def edit_cell(path: Path,sheet_name: str,coordinate: str,*,value=None,remove_formula: bool=False) -> None:
    with zipfile.ZipFile(path,"r") as archive: target_sheet=sheet_xml_path(archive,sheet_name)
    changed=False; namespace="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    def transform(name: str,data: bytes) -> bytes:
        nonlocal changed
        if name!=target_sheet: return data
        root=etree.fromstring(data); cells=root.xpath(f'//m:c[@r="{coordinate}"]',namespaces={"m":namespace})
        if not cells: raise ValueError(f"{sheet_name}!{coordinate} 不存在")
        cell=cells[0]
        if remove_formula:
            node=cell.find(f"{{{namespace}}}f")
            if node is None: raise ValueError(f"{sheet_name}!{coordinate} 没有公式")
            cell.remove(node)
        if value is not None:
            inline=cell.find(f"{{{namespace}}}is")
            if inline is not None: cell.remove(inline)
            node=cell.find(f"{{{namespace}}}v")
            if node is None: node=etree.SubElement(cell,f"{{{namespace}}}v")
            if isinstance(value,bool): node.text="1" if value else "0"; cell.set("t","b")
            elif isinstance(value,(int,float)): node.text=str(value); cell.attrib.pop("t",None)
            else: node.text=str(value); cell.set("t","str")
        changed=True; return etree.tostring(root,xml_declaration=True,encoding="UTF-8",standalone=True)
    rewrite_zip(path,transform)
    if not changed: raise ValueError("目标单元格未修改")


def openpyxl_mutation(path: Path,mutator: Callable) -> None:
    workbook=load_workbook(path); mutator(workbook); workbook.save(path)
def mutate_json(path: Path,mutator: Callable[[dict],None]) -> None:
    payload=json.loads(path.read_text(encoding="utf-8")); mutator(payload); path.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
def decode_message(message) -> str:
    if message.is_multipart():
        return "\n".join((part.get_payload(decode=True) or b"").decode(part.get_content_charset() or "utf-8",errors="replace") for part in message.walk() if part.get_content_type()=="text/plain").strip()
    return (message.get_payload(decode=True) or b"").decode(message.get_content_charset() or "utf-8",errors="replace").strip()
def mbox_states(path: Path) -> list[dict]:
    box=mailbox.mbox(path,create=False); messages=list(box); box.close(); headers=(("from","From"),("to","To"),("cc","Cc"),("bcc","Bcc"),("subject","Subject"),("date","Date"),("message_id","Message-ID"),("status","X-EvoCowork-Status"),("order","X-EvoCowork-Order"),("source_row","X-EvoCowork-Source-Row"),("decision","X-EvoCowork-Decision"),("action","X-EvoCowork-Action"),("candidate_outreach","X-EvoCowork-Candidate-Outreach-Executed"),("source_modified","X-EvoCowork-Source-Modified"))
    return [{**{key:str(message.get(header,"")) for key,header in headers},"body":decode_message(message)} for message in messages]
def write_mbox_states(path: Path,states: list[dict]) -> None:
    chunks=[]; headers=(("from","From"),("to","To"),("cc","Cc"),("bcc","Bcc"),("subject","Subject"),("date","Date"),("message_id","Message-ID"),("status","X-EvoCowork-Status"),("order","X-EvoCowork-Order"),("source_row","X-EvoCowork-Source-Row"),("decision","X-EvoCowork-Decision"),("action","X-EvoCowork-Action"),("candidate_outreach","X-EvoCowork-Candidate-Outreach-Executed"),("source_modified","X-EvoCowork-Source-Modified"))
    for order,state in enumerate(states,1):
        message=EmailMessage(policy=SMTP)
        for key,header in headers:
            if state.get(key): message[header]=state[key]
        message.set_content(state["body"].rstrip()+"\n",charset="utf-8"); buffer=BytesIO(); BytesGenerator(buffer,policy=SMTP,mangle_from_=True,maxheaderlen=0).flatten(message); chunks.extend([f"From MAILER-DAEMON Wed Apr 15 09:00:{order:02d} 2026\n".encode("ascii"),buffer.getvalue().replace(b"\r\n",b"\n"),b"\n"])
    path.write_bytes(b"".join(chunks))
def mutate_message(path: Path,index: int,mutator: Callable[[dict],None]) -> None:
    states=mbox_states(path); mutator(states[index]); write_mbox_states(path,states)


def replace_docx_text(path: Path,old: str,new: str,*,row_token: str|None=None,min_cells: int=0) -> None:
    document=Document(path); changed=False
    if row_token is None:
        for paragraph in document.paragraphs:
            if old in paragraph.text: paragraph.text=paragraph.text.replace(old,new,1); changed=True; break
    if not changed:
        for table in document.tables:
            for row in table.rows:
                if len(row.cells)<min_cells: continue
                if row_token and row_token not in " | ".join(cell.text for cell in row.cells): continue
                for cell in row.cells:
                    if old in cell.text: cell.text=cell.text.replace(old,new,1); changed=True; break
                if changed: break
            if changed: break
    if not changed: raise ValueError(f"Word 找不到 {old!r}")
    document.save(path)
def replace_pptx_text(path: Path,old: str,new: str,*,row_token: str|None=None,min_cells: int=0) -> None:
    presentation=Presentation(path); changed=False
    for slide in presentation.slides:
        for shape in slide.shapes:
            if row_token is None and getattr(shape,"has_text_frame",False) and old in shape.text: shape.text=shape.text.replace(old,new,1); changed=True; break
            if getattr(shape,"has_table",False):
                for row in shape.table.rows:
                    if len(row.cells)<min_cells: continue
                    if row_token and row_token not in " | ".join(cell.text for cell in row.cells): continue
                    for cell in row.cells:
                        if old in cell.text: cell.text=cell.text.replace(old,new,1); changed=True; break
                    if changed: break
            if changed: break
        if changed: break
    if not changed: raise ValueError(f"PPT 找不到 {old!r}")
    presentation.save(path)
def drop_ppt_slide(path: Path) -> None:
    presentation=Presentation(path); presentation.slides._sldIdLst.remove(presentation.slides._sldIdLst[-1]); presentation.save(path)


def build_cases() -> list[Case]:
    cases: list[Case]=[]
    for name in FILES: cases.append((f"missing-{name}",lambda root,n=name:(root/n).unlink(),"R001"))
    cases.append(("extra-file",lambda root:(root/"unexpected.txt").write_text("unexpected",encoding="utf-8"),"R001"))
    corrupt={XLSX:"R001",DOCX:"R001",PPTX:"R001",MBOX:"R011",ACTIONS:"R012",SLACK:"R013",EXECUTION:"R013"}
    for name in FILES: cases.append((f"corrupt-{name}",lambda root,n=name:(root/n).write_bytes(b"not-a-valid-artifact"),corrupt[name]))

    source_wrong={"A":99,"B":"Wrong Role","C":"Wrong Dept","D":"Wrong Location","E":"Wrong Manager","F":46000,"G":"Wrong Status","H":"Wrong Notes"}
    for row in range(2,10):
        for column,value in source_wrong.items(): cases.append((f"source-{column}{row}",lambda root,c=f"{column}{row}",v=value:edit_cell(root/XLSX,"Source_Positions",c,value=v),"R002"))
    for row in range(2,4):
        for column,value in (("A","WRONG-ID"),("B","WRONG-THREAD"),("C","wrong@company.example.com"),("D","wrong@company.example.com"),("E","Wrong subject"),("F","Wrong date"),("G","WRONG"),("H","Wrong body")):
            cases.append((f"source-message-{column}{row}",lambda root,c=f"{column}{row}",v=value:edit_cell(root/XLSX,"Source_Messages",c,value=v),"R003"))
    for row in range(2,9):
        for column,value in (("A","WRONG"),("B","WRONG"),("C","Wrong requirement"),("D","Wrong authority"),("E","wrong-source")): cases.append((f"policy-{column}{row}",lambda root,c=f"{column}{row}",v=value:edit_cell(root/XLSX,"Policy_Source",c,value=v),"R003"))
    for row in range(2,22):
        for column,value in (("A","WRONG-ID"),("B","WRONG-TYPE"),("C","Wrong subject"),("D","Wrong fact"),("E","wrong-source")): cases.append((f"evidence-{column}{row}",lambda root,c=f"{column}{row}",v=value:edit_cell(root/XLSX,"Evidence_Register",c,value=v),"R003"))

    for row in range(2,10):
        for column in range(1,29):
            coordinate=f"{get_column_letter(column)}{row}"; cases.append((f"formula-{coordinate}",lambda root,c=coordinate:edit_cell(root/XLSX,"Posting_Decisions",c,remove_formula=True),"R004"))
    row_criteria={2:"R005",3:"R008",4:"R006",5:"R005",6:"R008",7:"R005",8:"R007",9:"R008"}
    for row,criterion in row_criteria.items():
        for column in range(10,29):
            coordinate=f"{get_column_letter(column)}{row}"; cases.append((f"decision-{coordinate}",lambda root,c=coordinate:edit_cell(root/XLSX,"Posting_Decisions",c,remove_formula=True,value="WRONG"),criterion))
    direct_criteria={2:"R005",3:"R006",4:"R005",5:"R005"}
    for row,criterion in direct_criteria.items():
        for column in range(1,18):
            coordinate=f"{get_column_letter(column)}{row}"; cases.append((f"direct-{coordinate}",lambda root,c=coordinate:edit_cell(root/XLSX,"Direct_Repost_Queue",c,value="WRONG"),criterion))
    for column in range(1,17):
        coordinate=f"{get_column_letter(column)}2"; cases.append((f"stale-{coordinate}",lambda root,c=coordinate:edit_cell(root/XLSX,"Stale_Approval_Queue",c,value="WRONG"),"R007"))
    for row in range(2,5):
        for column in range(1,14):
            coordinate=f"{get_column_letter(column)}{row}"; cases.append((f"no-action-{coordinate}",lambda root,c=coordinate:edit_cell(root/XLSX,"No_Action_Control",c,value="WRONG"),"R008"))
    communication_criteria={2:"R005",3:"R006",4:"R005",5:"R005",6:"R007",7:"R009"}
    for row,criterion in communication_criteria.items():
        for column in range(1,13):
            coordinate=f"{get_column_letter(column)}{row}"; cases.append((f"communication-{coordinate}",lambda root,c=coordinate:edit_cell(root/XLSX,"Communication_Register",c,value="WRONG"),criterion))
    for column in range(1,13):
        coordinate=f"{get_column_letter(column)}2"; cases.append((f"cold-{coordinate}",lambda root,c=coordinate:edit_cell(root/XLSX,"Cold_Outreach_Control",c,value="WRONG"),"R009"))
    for row in range(2,22): cases.append((f"summary-B{row}",lambda root,c=f"B{row}":edit_cell(root/XLSX,"Control_Summary",c,value=99),"R010"))
    cases.extend([("xlsx-table-removed",lambda root:openpyxl_mutation(root/XLSX,lambda wb:wb["Source_Positions"].tables.clear()),"R017"),("xlsx-chart-removed",lambda root:openpyxl_mutation(root/XLSX,lambda wb:setattr(wb["Control_Summary"],"_charts",[])),"R017"),("xlsx-cf-removed",lambda root:openpyxl_mutation(root/XLSX,lambda wb:setattr(wb["Posting_Decisions"],"conditional_formatting",type(wb["Posting_Decisions"].conditional_formatting)())),"R017")])

    names=["Frontend Engineer","Account Executive","DevOps Engineer","Product Designer","Senior Backend Engineer","COLD_OUTREACH_ESCALATION"]; criteria=["R005","R006","R005","R005","R007","R009"]
    for index,(name,criterion) in enumerate(zip(names,criteria)):
        specific=(("from","wrong@company.example.com",criterion),("to","wrong@company.example.com",criterion),("cc","cc@company.example.com",criterion),("bcc","bcc@company.example.com",criterion),("subject","Wrong subject",criterion),("body","Wrong body",criterion),("date","Thu, 16 Apr 2026 09:00:00 +0000","R011"),("message_id","<wrong@evocowork.local>","R011"),("status","SENT","R011"),("order","99","R011"),("source_row","99","R011"),("decision","WRONG","R011"),("action","WRONG","R011"),("candidate_outreach","YES","R011"),("source_modified","YES","R011"))
        for field,value,expected in specific: cases.append((f"mbox-{index}-{field}",lambda root,i=index,f=field,v=value:mutate_message(root/MBOX,i,lambda state:state.update({f:v})),expected))
    cases.extend([
        ("mbox-talent-add-senior",lambda root:mutate_message(root/MBOX,5,lambda state:state.update({"body":state["body"]+"\nSenior Backend Engineer"})),"R007"),
        ("mbox-talent-add-marketing",lambda root:mutate_message(root/MBOX,5,lambda state:state.update({"body":state["body"]+"\nMarketing Manager"})),"R008"),
        ("mbox-candidate-invented",lambda root:mutate_message(root/MBOX,5,lambda state:state.update({"body":state["body"]+"\nContact Alex Candidate at alex@example.com."})),"R009"),
    ])

    global_action_mutations=(("task_id","wrong"),("status","EXECUTED"),("frozen_time","2026-04-16T09:00:00Z"),("stale_rule","at least 90 days"),("position_count",7),("direct_repost_count",5),("stale_vp_approval_count",0),("no_action_count",2),("email_draft_count",5),("emails_sent",1),("postings_reposted",1),("source_updates_applied",1),("candidate_outreach_executed",1),("slack_posts",1))
    for key,value in global_action_mutations: cases.append((f"actions-{key}",lambda root,k=key,v=value:mutate_json(root/ACTIONS,lambda payload:payload.update({k:v})),"R012"))
    cases.extend([("actions-precedence",lambda root:mutate_json(root/ACTIONS,lambda payload:payload["precedence"].reverse()),"R012"),("actions-drop-first",lambda root:mutate_json(root/ACTIONS,lambda payload:payload["decisions"].pop(0)),"R012"),("actions-reverse",lambda root:mutate_json(root/ACTIONS,lambda payload:payload["decisions"].reverse()),"R012")])
    entry_wrong={"source_row":99,"title":"Wrong Role","department":"Wrong Dept","location":"Wrong Location","hiring_manager":"Wrong Manager","original_post_date":"2027-01-01","age_days":999,"source_status":"Wrong","notes":"Wrong","decision":"WRONG","communication_type":"WRONG","recipient":"wrong@company.example.com","email_subject":"Wrong","proposed_post_date":"2027-01-01","source_update_applied":True,"posting_reposted":True,"candidate_outreach_executed":True,"slack_posted":True,"state":"EXECUTED","source_trace":[]}
    entry_criteria=["R005","R008","R006","R005","R008","R005","R007","R008"]
    for index,criterion in enumerate(entry_criteria):
        for key,value in entry_wrong.items():
            expected="R009" if key=="candidate_outreach_executed" else criterion
            cases.append((f"actions-{index}-{key}",lambda root,i=index,k=key,v=value:mutate_json(root/ACTIONS,lambda payload:payload["decisions"][i].update({k:v})),expected))

    for key,value in (("task_id","wrong"),("decision","AUTHORIZED"),("reason","Publish now"),("draft_count",1),("posted_count",1),("posted",True),("job_titles_in_slack",["Frontend Engineer"]),("candidate_information_in_slack",["Alex"]),("messages_created",["posted"]),("source_trace",[])):
        cases.append((f"slack-{key}",lambda root,k=key,v=value:mutate_json(root/SLACK,lambda payload:payload.update({k:v})),"R013"))
    cases.extend([("slack-channel",lambda root:mutate_json(root/SLACK,lambda payload:payload["source_channel"].update({"id":"C_OTHER"})),"R013"),("slack-private",lambda root:mutate_json(root/SLACK,lambda payload:payload["source_channel"].update({"is_private":True})),"R013")])

    execution_mutations=[
        ("task-id","R013",lambda p:p.update({"task_id":"wrong"})),("status","R013",lambda p:p.update({"status":"EXECUTED"})),("side-effect","R013",lambda p:p.update({"real_side_effect_count":1})),
        ("email-direct","R013",lambda p:p["email"].update({"direct_repost_drafts":3})),("email-vp","R013",lambda p:p["email"].update({"stale_vp_approval_drafts":0})),("email-talent","R013",lambda p:p["email"].update({"talent_lead_escalation_drafts":0})),("email-total","R013",lambda p:p["email"].update({"total_drafts":5})),("email-sent","R013",lambda p:p["email"].update({"sent_count":1,"state":"SENT"})),
        ("posting-proposal","R007",lambda p:p["posting"].update({"direct_repost_proposals":5})),("posting-reposted","R007",lambda p:p["posting"].update({"actually_reposted":1,"state":"REPOSTED"})),
        ("tracker-draft","R013",lambda p:p["source_tracker"].update({"update_drafts":1})),("tracker-applied","R013",lambda p:p["source_tracker"].update({"updates_applied":1,"source_modified":True,"state":"MODIFIED"})),
        ("cold-policy","R009",lambda p:p["cold_outreach"].update({"policy_decision":"EXECUTE"})),("cold-lead","R009",lambda p:p["cold_outreach"].update({"talent_lead":"wrong@company.example.com"})),("cold-reconsent","R009",lambda p:p["cold_outreach"].update({"candidate_reconsent_required":False})),("cold-records","R009",lambda p:p["cold_outreach"].update({"candidate_records_in_input":1})),("cold-drafts","R009",lambda p:p["cold_outreach"].update({"candidate_drafts_created":1})),("cold-sent","R009",lambda p:p["cold_outreach"].update({"candidate_messages_sent":1,"executed":True})),
        ("slack-draft","R013",lambda p:p["slack"].update({"draft_count":1})),("slack-post","R013",lambda p:p["slack"].update({"posted_count":1,"state":"POSTED"})),("trace","R013",lambda p:p.update({"source_trace":[]})),
    ]
    for name,criterion,mutator in execution_mutations: cases.append((f"execution-{name}",lambda root,m=mutator:mutate_json(root/EXECUTION,m),criterion))

    word_cases=[
        ("summary-direct","4","5","DIRECT_REPOST","R015",0),("summary-stale","1","0","STALE_VP_APPROVAL","R015",0),("summary-freeze","1","0","FREEZE_NO_REPOST","R015",0),("summary-hold","1","0","HOLD_NO_REPOST","R015",0),("summary-filled","1","0","FILLED_NO_ACTION","R015",0),
        ("frontend-decision","DIRECT_REPOST","FREEZE_NO_REPOST","Frontend Engineer","R005",11),("marketing-decision","FREEZE_NO_REPOST","DIRECT_REPOST","Marketing Manager","R008",11),("account-age","90","91","Account Executive","R006",11),("devops-route","recruiting@company.example.com","wrong@company.example.com","DevOps Engineer","R005",11),("customer-decision","FILLED_NO_ACTION","DIRECT_REPOST","Customer Success Manager","R008",11),("product-route","recruiting@company.example.com","wrong@company.example.com","Product Designer","R005",11),("senior-route","vp-people@company.example.com","recruiting@company.example.com","Senior Backend Engineer","R007",11),("sdr-decision","HOLD_NO_REPOST","DIRECT_REPOST","Sales Development Rep","R008",11),
        ("candidate-state","PROHIBITED","CONTACTED","Candidate outreach","R013",0),("slack-state","NOT_AUTHORIZED_CHANNEL","POSTED","Slack","R013",0),("email-state","DRAFT_NOT_SENT","SENT","Emails","R013",0),("posting-state","NOT_REPOSTED","REPOSTED","Postings","R013",0),
    ]
    for name,old,new,row_token,criterion,min_cells in word_cases: cases.append((f"word-{name}",lambda root,o=old,n=new,t=row_token,m=min_cells:replace_docx_text(root/DOCX,o,n,row_token=t,min_cells=m),criterion))
    narrative_cases=[("only roles older than 90 days require VP approval, so the Account Executive at exactly 90 days remains a direct repost","Account Executive at 90 days requires VP approval","R006"),("directly reposted","requires VP approval","R006"),("Senior Backend Engineer is 156 days","Senior Backend Engineer is 90 days","R007"),("not reposted before approval","already reposted","R007"),("Marketing Manager is frozen","Marketing Manager is reposted","R008"),("Customer Success is Filled","Customer Success is Open","R008"),("Sales Development Rep is on hold","Sales Development Rep is reposted","R008"),("Talent Lead for re-consent governance","HR Ops executes outreach","R009"),("no candidate identity is invented or contacted","candidate identity was invented","R009"),("Posted: 2026-04-15","Posted: 2026-04-16","R015"),("public #recruiting channel remains at zero drafts and zero posts","#recruiting has one post","R015")]
    for old,new,criterion in narrative_cases: cases.append((f"word-narrative-{old[:12]}",lambda root,o=old,n=new:replace_docx_text(root/DOCX,o,n),criterion))

    ppt_cases=[("summary-direct","4","5","Frontend; Account; DevOps; Product","R014"),("summary-stale","1","0","Senior Backend","R014"),("summary-freeze","1","0","Marketing Manager","R014"),("summary-hold","1","0","Sales Development Rep","R014"),("summary-filled","1","0","Customer Success","R014"),("boundary","only more than 90 days routes to VP People","90 days routes to VP People","","R006"),("cold","Talent Lead + re-consent required","HR Ops contacts candidates","","R009"),("email-count","6 drafts · 0 sent","5 drafts · 1 sent","","R013"),("post-count","4 proposed · 0 reposted","4 proposed · 1 reposted","","R013"),("candidate-count","0 contacted · 0 posts","1 contacted · 1 post","","R013")]
    for name,old,new,row_token,criterion in ppt_cases: cases.append((f"ppt-{name}",lambda root,o=old,n=new,t=row_token or None:replace_pptx_text(root/PPTX,o,n,row_token=t),criterion))
    ppt_rows=[("frontend","Frontend Engineer","DIRECT_REPOST","FREEZE_NO_REPOST","R005"),("marketing","Marketing Manager","FREEZE_NO_REPOST","DIRECT_REPOST","R008"),("account","Account Executive","90","91","R006"),("devops","DevOps Engineer","recruiting@company.example.com","wrong@company.example.com","R005"),("customer","Customer Success Manager","FILLED_NO_ACTION","DIRECT_REPOST","R008"),("product","Product Designer","DIRECT_REPOST","HOLD_NO_REPOST","R005"),("senior","Senior Backend Engineer","vp-people@company.example.com","recruiting@company.example.com","R007"),("sdr","Sales Development Rep","HOLD_NO_REPOST","DIRECT_REPOST","R008"),("talent","COLD_OUTREACH_ESCALATION","talent-lead@company.example.com","wrong@company.example.com","R009")]
    for name,token,old,new,criterion in ppt_rows:
        minimum_cells=0 if name=="talent" else 10
        cases.append((f"ppt-row-{name}",lambda root,t=token,o=old,n=new,m=minimum_cells:replace_pptx_text(root/PPTX,o,n,row_token=t,min_cells=m),criterion))
    cases.append(("ppt-drop-slide",lambda root:drop_ppt_slide(root/PPTX),"R014"))
    return cases


def main() -> None:
    parser=argparse.ArgumentParser(); parser.add_argument("--json-out",type=Path); args=parser.parse_args(); judge=load_judge(); cases=build_cases(); results=[]; failures=[]
    for name,mutator,expected_failure in cases:
        with tempfile.TemporaryDirectory(prefix=f"negative-5123-{name}-") as temporary:
            artifact_root=Path(temporary)/"artifacts"; shutil.copytree(GOLDEN,artifact_root)
            try:
                mutator(artifact_root); result=judge.evaluate(artifact_root); detected=not result.get("criteria",{}).get(expected_failure,{}).get("passed",True); rejected=not result.get("passed",False) and float(result.get("reward",1.0))<0.8
                results.append({"case":name,"expected_failure":expected_failure,"detected":detected,"rejected":rejected,"reward":result.get("reward"),"hurdle_failures":result.get("hurdle_failures",[])})
                if not detected or not rejected: failures.append(f"{name}: expected={expected_failure}, detected={detected}, reward={result.get('reward')}")
            except Exception as exc: failures.append(f"{name}: mutation/evaluation error: {type(exc).__name__}: {exc}")
    payload={"task_id":"cowork-evo/automationbench-5123-job-posting-governance-package","case_count":len(results),"declared_case_count":len(cases),"all_rejected":not failures and len(results)==len(cases),"max_reward":max((float(item["reward"]) for item in results),default=None),"failures":failures,"results":results}
    if args.json_out: args.json_out.parent.mkdir(parents=True,exist_ok=True); args.json_out.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(payload,ensure_ascii=False,indent=2))
    if failures or len(results)!=len(cases): raise SystemExit(1)


if __name__=="__main__": main()
