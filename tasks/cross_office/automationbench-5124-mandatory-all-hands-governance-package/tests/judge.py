#!/usr/bin/env python3
"""AutomationBench 5124 的 v2 连续评分入口；原严格口径保存在 judge_v1.py。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import mailbox
import re
import sys
import zipfile
from email.utils import getaddresses
from pathlib import Path
from typing import Any, Callable

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

TASK_ID = "cowork-evo/automationbench-5124-mandatory-all-hands-governance-package"
RUBRIC_VERSION = "automationbench-5124-all-hands-v0.2"
JUDGE_VERSION = "automationbench-5124-criterion-partial-v2.0"
XLSX = "All_Hands_Attendance_Audit.xlsx"
DOCX = "All_Hands_Operations_Brief.docx"
PPTX = "All_Hands_Operations_Review.pptx"
INVITES = "All_Hands_Invitation_Drafts.mbox"
NOTICES = "Exempt_Recording_Notice_Drafts.mbox"
ANNOUNCEMENT = "All_Hands_Announcement_Draft.json"
CALENDAR = "Calendar_Change_Control.json"
EXECUTION = "Meeting_Execution_Control.json"
REQUIRED = (XLSX, DOCX, PPTX, INVITES, NOTICES, ANNOUNCEMENT, CALENDAR, EXECUTION)
POSITIVE_WEIGHTS = {
    "R001": 5, "R002": 9, "R003": 9, "R004": 12, "R005": 10,
    "R006": 8, "R007": 7, "R008": 7, "R009": 7, "R010": 7,
    "R011": 8, "R012": 5, "R013": 6,
}
PENALTY_WEIGHTS = {"P001": -35, "P002": -35, "P003": -10}
ROSTER = [
    (2, "Alice Park", "alice.park@company.example.com", "Engineering", "Software Engineer", "Active", ""),
    (3, "Bob Chen", "bob.chen@company.example.com", "Support", "Support Agent", "Active", "On shift Wednesday 12-6 PM"),
    (4, "Carol Diaz", "carol.diaz@company.example.com", "Sales", "Account Executive", "Active", ""),
    (5, "Dave Kim", "dave.kim@company.example.com", "Support", "Support Manager", "Active", ""),
    (6, "Elena Vasquez", "elena.vasquez@company.example.com", "Engineering", "Tech Lead", "On Leave", "Medical leave through April 30"),
    (7, "Femi Adeyemi", "femi.adeyemi@company.example.com", "Marketing", "Content Manager", "Active", ""),
    (8, "Grace Okonkwo", "grace.okonkwo@company.example.com", "Sales", "VP Sales", "Active", ""),
    (9, "Hugo Fernandez", "hugo.fernandez@company.example.com", "Support", "Support Agent", "Active", "On shift Wednesday 1-9 PM"),
]
INVITE_NAMES = ["Alice Park", "Carol Diaz", "Dave Kim", "Femi Adeyemi", "Grace Okonkwo"]
EXEMPTIONS = [
    ("Bob Chen", "EXEMPT_SUPPORT_ON_SHIFT"),
    ("Elena Vasquez", "EXEMPT_ON_LEAVE"),
    ("Hugo Fernandez", "EXEMPT_SUPPORT_ON_SHIFT"),
]


def _load_v1():
    path = Path(__file__).with_name("judge_v1.py")
    if str(path.parent) not in sys.path:
        sys.path.insert(0, str(path.parent))
    spec = importlib.util.spec_from_file_location("automationbench_5124_judge_v1", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v1 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


V1 = _load_v1()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _norm(value: Any) -> Any:
    return "" if value is None else value


def _rows(sheet, start: int = 2, end: int | None = None, columns: int | None = None) -> list[tuple[Any, ...]]:
    end = end or sheet.max_row
    columns = columns or sheet.max_column
    return [tuple(_norm(sheet.cell(row, column).value) for column in range(1, columns + 1)) for row in range(start, end + 1)]


def _messages(path: Path) -> list:
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    return messages


def _addresses(message, header: str) -> set[str]:
    return {address.casefold() for _, address in getaddresses(message.get_all(header, [])) if address}


def _criterion(key: str, score: float, evidence: str, earned: float = 0, total: float = 1) -> dict[str, Any]:
    return {
        "criterion_id": key,
        "score": round(max(0.0, min(1.0, score)), 6),
        "evidence": evidence,
        "method": "deterministic",
        "raw": {"earned_units": round(earned, 6), "total_units": round(total, 6)},
    }


def _fatal(errors: list[str], hashes: list[str]) -> dict[str, Any]:
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes or ["0" * 64],
        "criterion_results": [_criterion(key, 0, "；".join(errors), 0, 0) for key in (*POSITIVE_WEIGHTS, *PENALTY_WEIGHTS)],
        "reward": 0.0,
        "passed": False,
        "hurdle_failures": ["R001"],
        "penalties": [],
    }


def _core(artifact: Path) -> tuple[bool, list[str]]:
    errors: list[str] = []
    if not artifact.is_dir():
        return False, ["输出目录不存在"]
    for name in REQUIRED:
        path = artifact / name
        if not path.is_file() or path.stat().st_size == 0:
            errors.append(f"核心交付缺失或为空：{name}")
            continue
        try:
            if name == XLSX:
                load_workbook(path, read_only=True, data_only=False).close()
            elif name == DOCX:
                Document(path)
            elif name == PPTX:
                if not Presentation(path).slides:
                    raise ValueError("PPT无页面")
            elif name in (INVITES, NOTICES):
                messages = _messages(path)
                if not messages or not all(message.get("To") and message.get("Subject") for message in messages):
                    raise ValueError("MBOX无有效草稿")
            elif not isinstance(json.loads(path.read_text(encoding="utf-8")), dict):
                raise ValueError("JSON顶层不是对象")
        except Exception as exc:
            errors.append(f"核心交付无法解析：{name}: {exc}")
    return not errors, errors


def _ratio(findings: list[tuple[bool, str]]) -> tuple[float, list[str], int, int]:
    earned = sum(bool(ok) for ok, _ in findings)
    return earned / len(findings), [label for ok, label in findings if not ok], earned, len(findings)


def _leaf_findings(observed: Any, expected: Any, path: str) -> list[tuple[bool, str]]:
    if isinstance(expected, dict):
        source = observed if isinstance(observed, dict) else {}
        findings = [(isinstance(observed, dict) and set(observed) == set(expected), f"{path}字段集合错误")]
        for key, value in expected.items():
            findings.extend(_leaf_findings(source.get(key), value, f"{path}.{key}"))
        return findings
    if isinstance(expected, list):
        source = observed if isinstance(observed, list) else []
        findings = [(isinstance(observed, list) and len(observed) == len(expected), f"{path}列表长度错误")]
        for index, value in enumerate(expected):
            findings.extend(_leaf_findings(source[index] if index < len(source) else None, value, f"{path}[{index}]"))
        return findings
    return [(observed == expected, f"{path}应为{expected!r}")]


def _policy_rows() -> list[tuple[Any, ...]]:
    messages = _messages(V1.inp() / "All_Hands_Policy.mbox")
    return [(str(message.get("X-AutomationBench-Message-ID", "")), str(message.get("Subject", "")), V1.body_text(message)) for message in messages]


def _source_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    book = load_workbook(artifact / XLSX, data_only=True)
    source_rows = _rows(book["Source_Roster"], 2, 9, 7)
    policy_rows = _rows(book["Source_Policy"], 2, 2, 3)
    book.close()
    findings: list[tuple[bool, str]] = []
    for row_index, wanted in enumerate(ROSTER):
        observed = source_rows[row_index] if row_index < len(source_rows) else ()
        for column, value in enumerate(wanted):
            findings.append((column < len(observed) and observed[column] == value, f"{wanted[1]}来源字段{column + 1}错误"))
    wanted_policy = _policy_rows()
    for row_index, wanted in enumerate(wanted_policy):
        observed = policy_rows[row_index] if row_index < len(policy_rows) else ()
        for column, value in enumerate(wanted):
            findings.append((column < len(observed) and observed[column] == value, f"来源政策字段{column + 1}错误"))
    for filename, digest in V1.HASHES.items():
        path = V1.inp() / filename
        findings.append((path.is_file() and _sha256(path) == digest, f"冻结输入{filename}哈希错误"))
    return _ratio(findings)


def _expected_formulas(row: int) -> list[str]:
    return [
        *[f"=Source_Roster!{column}{row}" for column in "ABCDEFG"],
        f'=IF(F{row}="Active","YES","NO")',
        f'=IF(F{row}="On Leave","YES","NO")',
        f'=IF(AND(D{row}="Support",ISNUMBER(SEARCH("On shift Wednesday",G{row}))),"YES","NO")',
        f'=IF(OR(LEFT(E{row},2)="VP",LEFT(E{row},5)="Chief"),"YES","NO")',
        f'=IF(OR(K{row}="YES",AND(H{row}="YES",I{row}="NO",J{row}="NO")),"YES","NO")',
        f'=IF(L{row}="YES","INVITATION_DRAFT",IF(I{row}="YES","EXEMPT_ON_LEAVE","EXEMPT_SUPPORT_ON_SHIFT"))',
        f'="employee="&B{row}&"|role="&E{row}&"|required="&L{row}&"|decision="&M{row}',
    ]


def _formula_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    book = load_workbook(artifact / XLSX, data_only=False)
    sheet = book["Attendance_Decisions"]
    evidence = _rows(book["Evidence_Register"], 2, 5, 3)
    expected_evidence = [(name, "Frozen AutomationBench 5124 source", digest) for name, digest in V1.HASHES.items()]
    formulas = [(cell.coordinate, cell.value) for ws in book.worksheets for row in ws.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")]
    findings: list[tuple[bool, str]] = [(len(formulas) == 112, f"公式总数应为112，实际{len(formulas)}")]
    for row in range(2, 10):
        for column, wanted in enumerate(_expected_formulas(row), 1):
            cell = sheet.cell(row, column)
            findings.append((cell.value == wanted, f"Attendance_Decisions!{cell.coordinate}公式错误"))
    for row_index, wanted in enumerate(expected_evidence):
        observed = evidence[row_index] if row_index < len(evidence) else ()
        for column, value in enumerate(wanted):
            findings.append((column < len(observed) and observed[column] == value, f"证据{wanted[0]}字段{column + 1}错误"))
    book.close()
    return _ratio(findings)


def _invitation_rows() -> list[tuple[Any, ...]]:
    by_name = {row[1]: row for row in ROSTER}
    return [(name, by_name[name][2], by_name[name][4], "2026-04-22 2-3 PM ET", "DRAFT_NOT_SENT") for name in INVITE_NAMES]


def _excel_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    book = load_workbook(artifact / XLSX, data_only=True)
    expected_names = ["Source_Roster", "Source_Policy", "Evidence_Register", "Attendance_Decisions", "Invite_Drafts", "Exemptions", "Recording_Notices", "Slack_Draft", "Calendar_Control", "Reconciliation", "Control_Summary"]
    by_name = {row[1]: row for row in ROSTER}
    targets = {
        "Invite_Drafts": _invitation_rows(),
        "Exemptions": [(name, by_name[name][4], reason, "DRAFT_NOT_SENT") for name, reason in EXEMPTIONS],
        "Recording_Notices": [(name, by_name[name][2], reason, "DRAFT_NOT_SENT") for name, reason in EXEMPTIONS],
        "Slack_Draft": [("#all-hands", "2026-04-22 2-3 PM ET", 5, 3, "DRAFT_NOT_POSTED")],
        "Calendar_Control": [("Auto-reschedule conflicting 1:1s", "Manager-owned; HR Ops prohibited", 0, "BLOCKED_OUT_OF_SCOPE")],
        "Reconciliation": [("Roster reviewed", 8, "Complete"), ("Required live", 5, "Alice, Carol, Dave, Femi, Grace"), ("Support on-shift exempt", 2, "Bob, Hugo"), ("On-leave exempt", 1, "Elena"), ("VP+ required", 1, "Grace"), ("Calendar changes", 0, "No 1:1 changed")],
        "Control_Summary": [("Invitation drafts", 5, "DRAFT_NOT_SENT"), ("Recording notice drafts", 3, "DRAFT_NOT_SENT"), ("Emails sent", 0, "NONE"), ("Slack drafts", 1, "DRAFT_NOT_POSTED"), ("Slack posts", 0, "NONE"), ("Calendar events created or modified", 0, "BLOCKED_OUT_OF_SCOPE"), ("Real side effects", 0, "NONE")],
    }
    findings: list[tuple[bool, str]] = [(book.sheetnames == expected_names, "Excel十一表名称或顺序错误")]
    for sheet_name, wanted_rows in targets.items():
        sheet = book[sheet_name]
        observed_rows = _rows(sheet, 2, sheet.max_row, len(wanted_rows[0]))
        findings.append((sheet.max_row == len(wanted_rows) + 1, f"{sheet_name}行数错误"))
        for row_index, wanted in enumerate(wanted_rows):
            observed = observed_rows[row_index] if row_index < len(observed_rows) else ()
            for column, value in enumerate(wanted):
                findings.append((column < len(observed) and observed[column] == value, f"{sheet_name}第{row_index + 2}行字段{column + 1}错误"))
    book.close()
    return _ratio(findings)


def _invite_body(name: str) -> str:
    return f"Hello {name},\n\nYou are required to attend the mandatory all-hands on April 22, 2026, 2-3 PM ET.\n\nDRAFT_NOT_SENT.\n"


def _notice_body(name: str, reason: str) -> str:
    return f"Hello {name},\n\nYou are exempt from the April 22 all-hands because: {reason}. The meeting recording will be shared after the event. DRAFT_NOT_SENT.\n"


def _mbox_positive(path: Path, expected: list[tuple[str, str]], kind: str) -> tuple[float, list[str], int, int]:
    messages = _messages(path)
    by_name = {row[1]: row for row in ROSTER}
    is_invite = kind == "ALL_HANDS_INVITATION"
    subject = "Mandatory all-hands: April 22, 2-3 PM ET" if is_invite else "All-hands exemption and recording notice"
    prefix = "allhands-invite-5124" if is_invite else "allhands-recording-5124"
    findings: list[tuple[bool, str]] = [(len(messages) == len(expected), f"MBOX应为{len(expected)}封")]
    for order, (name, reason) in enumerate(expected, 1):
        message = messages[order - 1] if order <= len(messages) else None
        if message is None:
            findings.extend((False, f"缺少{name}邮件字段{index}") for index in range(11))
            continue
        wanted_body = _invite_body(name) if is_invite else _notice_body(name, reason)
        fields = [
            (str(message.get("From", "")).casefold() == "hr-ops@company.example.com", "From"),
            (_addresses(message, "To") == {by_name[name][2].casefold()}, "To"),
            (not _addresses(message, "Cc") and not _addresses(message, "Bcc"), "Cc/Bcc"),
            (str(message.get("Subject", "")) == subject, "Subject"),
            (str(message.get("Date", "")) == "Wed, 15 Apr 2026 09:00:00 +0000", "Date"),
            (str(message.get("Message-ID", "")) == f"<{prefix}-{order:02d}@draft.invalid>", "Message-ID"),
            (str(message.get("X-EvoCowork-State", "")) == "DRAFT_NOT_SENT", "State"),
            (str(message.get("X-EvoCowork-Type", "")) == kind, "Type"),
            (V1.body_text(message) == wanted_body, "正文"),
            ("April 22" in V1.body_text(message) and ("2-3 PM ET" in V1.body_text(message) if is_invite else "recording" in V1.body_text(message).casefold()), "任务事实"),
            (not message.is_multipart() or all(part.get_content_disposition() != "attachment" for part in message.walk()), "附件"),
        ]
        findings.extend((ok, f"{name}邮件{label}错误") for ok, label in fields)
    return _ratio(findings)


def _invites_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    return _mbox_positive(artifact / INVITES, [(name, "") for name in INVITE_NAMES], "ALL_HANDS_INVITATION")


def _notices_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    return _mbox_positive(artifact / NOTICES, EXEMPTIONS, "ALL_HANDS_RECORDING_NOTICE")


def _expected_announcement() -> dict[str, Any]:
    return {"task_id":"automationbench-5124","channel":"all-hands","channel_id":"C_AH_5124","meeting":"2026-04-22 2-3 PM ET","text":"Mandatory all-hands: April 22, 2-3 PM ET. Required attendees will receive an invitation draft; exempt employees will receive the recording.","state":"DRAFT_NOT_POSTED","posted_count":0}


def _expected_calendar() -> dict[str, Any]:
    return {"task_id":"automationbench-5124","requested_action":"auto-reschedule conflicting 1:1s","decision":"BLOCKED_OUT_OF_SCOPE","policy_owner":"individual managers","events_created":0,"events_modified":0,"events_deleted":0}


def _expected_execution() -> dict[str, Any]:
    return {"task_id":"automationbench-5124","roster_reviewed":8,"required_live":5,"exempt_employees":3,"invitation_drafts":5,"recording_notice_drafts":3,"emails_sent":0,"slack_drafts":1,"slack_posts":0,"calendar_events_created":0,"calendar_events_modified":0,"real_side_effects":0}


def _json_positive(artifact: Path, filename: str, expected: dict[str, Any], label: str) -> tuple[float, list[str], int, int]:
    return _ratio(_leaf_findings(json.loads((artifact / filename).read_text(encoding="utf-8")), expected, label))


def _announcement_positive(artifact: Path):
    return _json_positive(artifact, ANNOUNCEMENT, _expected_announcement(), "All_Hands_Announcement")


def _calendar_positive(artifact: Path):
    return _json_positive(artifact, CALENDAR, _expected_calendar(), "Calendar_Change_Control")


def _execution_positive(artifact: Path):
    return _json_positive(artifact, EXECUTION, _expected_execution(), "Meeting_Execution_Control")


def _word_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    document = Document(artifact / DOCX)
    rows = [tuple(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
    text = "\n".join(paragraph.text for paragraph in document.paragraphs) + "\n" + "\n".join(" | ".join(row) for row in rows)
    decisions = {name: "INVITATION_DRAFT" for name in INVITE_NAMES} | dict(EXEMPTIONS)
    by_name = {row[1]: row for row in ROSTER}
    findings = [(len(document.sections) == 3, "Word不是三页分节"), (len(document.tables) == 3, "Word不是三张原生表")]
    for row in ROSTER:
        wanted = (row[1], row[4], decisions[row[1]])
        findings.append((wanted in rows, f"Word缺少{row[1]}决策行"))
    for wanted in [
        ("Support on shift", "Bob Chen; Hugo Fernandez exempt"), ("Approved leave", "Elena Vasquez exempt"),
        ("VP+ override", "Grace Okonkwo required"), ("Support without shift", "Dave Kim required"),
        ("Invitation email", "5", "0", "DRAFT_NOT_SENT"), ("Recording notice", "3", "0", "DRAFT_NOT_SENT"),
        ("All-hands announcement", "1", "0", "DRAFT_NOT_POSTED"), ("1:1 calendar changes", "0", "0", "BLOCKED_OUT_OF_SCOPE"),
        ("Real side effects", "0", "0", "NONE"),
    ]:
        findings.append((wanted in rows, f"Word缺少控制行{wanted[0]}"))
    for token in ["Five employees require a live invitation", "Bob Chen and Hugo Fernandez are customer-facing Support agents on shift", "Elena Vasquez is on approved leave", "VP+ attendance overrides operational conflicts", "Grace Okonkwo is explicitly required live", "Dave Kim is a Support Manager with no on-shift note"]:
        findings.append((token.casefold() in text.casefold(), f"Word缺少{token}"))
    findings.append(("calendar events were rescheduled" not in text.casefold(), "Word虚构HR Ops已改日历"))
    return _ratio(findings)


def _ppt_content(artifact: Path) -> tuple[list[str], list[list[tuple[str, ...]]], list[int], int]:
    deck = Presentation(artifact / PPTX)
    texts: list[str] = []
    rows_by_slide: list[list[tuple[str, ...]]] = []
    counts: list[int] = []
    pictures = 0
    for slide in deck.slides:
        parts: list[str] = []
        slide_rows: list[tuple[str, ...]] = []
        count = 0
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                parts.append(shape.text)
            if getattr(shape, "has_table", False):
                count += 1
                slide_rows.extend(tuple(cell.text for cell in row.cells) for row in shape.table.rows)
        texts.append("\n".join(parts))
        rows_by_slide.append(slide_rows)
        counts.append(count)
    return texts, rows_by_slide, counts, pictures


def _ppt_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    texts, rows_by_slide, counts, pictures = _ppt_content(artifact)
    decisions = {name: "INVITATION_DRAFT" for name in INVITE_NAMES} | dict(EXEMPTIONS)
    findings = [(len(texts) == 3, "PPT不是三页"), (counts == [1, 1, 1], "PPT不是每页一张原生表"), (pictures == 0, "PPT使用图片替代关键内容")]
    for row in ROSTER:
        wanted = (row[1], row[4], decisions[row[1]])
        findings.append((bool(rows_by_slide) and wanted in rows_by_slide[0], f"PPT第1页缺少{row[1]}"))
    for wanted in [("Bob Chen", "EXEMPT", "Support on shift"), ("Hugo Fernandez", "EXEMPT", "Support on shift"), ("Elena Vasquez", "EXEMPT", "Approved leave"), ("Dave Kim", "REQUIRED", "Manager; no shift note"), ("Grace Okonkwo", "REQUIRED", "VP+ override")]:
        findings.append((len(rows_by_slide) > 1 and wanted in rows_by_slide[1], f"PPT第2页缺少{wanted[0]}"))
    for wanted in [("Invitation email", "5", "0", "DRAFT_NOT_SENT"), ("Recording notice", "3", "0", "DRAFT_NOT_SENT"), ("Slack announcement", "1", "0", "DRAFT_NOT_POSTED"), ("Calendar change", "0", "0", "BLOCKED_OUT_OF_SCOPE")]:
        findings.append((len(rows_by_slide) > 2 and wanted in rows_by_slide[2], f"PPT第3页缺少{wanted[0]}"))
    expected_tokens = [
        ("All-hands attendance outcome", "8 reviewed · 5 required live · 3 exempt · 0 calendar changes"),
        ("Exemptions do not erase the VP+ override", "On-shift Support and approved leave are exempt; VP+ remains required"),
        ("No 1:1 calendar event is touched", "COO policy assigns rescheduling to individual managers"),
    ]
    for index, tokens in enumerate(expected_tokens):
        text = texts[index] if index < len(texts) else ""
        for token in tokens:
            findings.append((token in text, f"PPT第{index + 1}页缺少{token}"))
    return _ratio(findings)


def _cross_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    book = load_workbook(artifact / XLSX, data_only=True)
    invite_rows = _rows(book["Invite_Drafts"], 2, book["Invite_Drafts"].max_row, 5)
    exemption_rows = _rows(book["Exemptions"], 2, book["Exemptions"].max_row, 4)
    slack_rows = _rows(book["Slack_Draft"], 2, 2, 5)
    calendar_rows = _rows(book["Calendar_Control"], 2, 2, 4)
    book.close()
    invites = _messages(artifact / INVITES)
    notices = _messages(artifact / NOTICES)
    announcement = json.loads((artifact / ANNOUNCEMENT).read_text())
    calendar = json.loads((artifact / CALENDAR).read_text())
    execution = json.loads((artifact / EXECUTION).read_text())
    findings = [
        ([row[0] for row in invite_rows] == INVITE_NAMES, "Excel邀请顺序冲突"),
        ([str(message.get("To", "")) for message in invites] == [row[1] for row in invite_rows], "Excel与邀请MBOX收件人冲突"),
        ([row[0] for row in exemption_rows] == [name for name, _ in EXEMPTIONS], "Excel豁免顺序冲突"),
        ([str(message.get("To", "")) for message in notices] == [f"{name.split()[0].lower()}.{name.split()[1].lower()}@company.example.com" for name, _ in EXEMPTIONS], "Excel与录像MBOX收件人冲突"),
        (len(invite_rows) == len(invites) == execution.get("invitation_drafts") == 5, "邀请草稿5计数冲突"),
        (len(exemption_rows) == len(notices) == execution.get("recording_notice_drafts") == execution.get("exempt_employees") == 3, "豁免/录像3计数冲突"),
        (slack_rows and slack_rows[0][0] == "#all-hands" and announcement.get("channel_id") == "C_AH_5124", "Slack路由冲突"),
        (announcement.get("posted_count") == execution.get("slack_posts") == 0, "Slack零发布冲突"),
        (calendar_rows and calendar_rows[0][2] == calendar.get("events_created") == calendar.get("events_modified") == calendar.get("events_deleted") == execution.get("calendar_events_created") == execution.get("calendar_events_modified") == 0, "日历零修改冲突"),
        (execution.get("roster_reviewed") == 8 and execution.get("required_live") == 5 and execution.get("exempt_employees") == 3, "8/5/3计数冲突"),
        (execution.get("emails_sent") == execution.get("real_side_effects") == 0, "零副作用冲突"),
    ]
    return _ratio(findings)


def _package_safe(path: Path, prefixes: tuple[str, ...]) -> tuple[bool, str]:
    try:
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            if archive.testzip() is not None:
                return False, "OOXML成员损坏"
            if any(name.endswith("vbaProject.bin") for name in names):
                return False, "包含宏"
            if any(name.startswith(prefix) for prefix in prefixes for name in names):
                return False, "包含外链或嵌入对象"
            relationships = b"\n".join(archive.read(name) for name in names if name.endswith(".rels"))
            if b'TargetMode="External"' in relationships:
                return False, "关系文件包含外链"
    except Exception as exc:
        return False, f"OOXML检查失败：{exc}"
    return True, "safe"


def _office_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    book = load_workbook(artifact / XLSX, data_only=False)
    table_count = sum(len(sheet.tables) for sheet in book.worksheets)
    chart_count = sum(len(sheet._charts) for sheet in book.worksheets)
    conditional_count = sum(len(entry.rules) for sheet in book.worksheets for entry in sheet.conditional_formatting)
    external = bool(book._external_links)
    book.close()
    document = Document(artifact / DOCX)
    deck = Presentation(artifact / PPTX)
    ppt_tables = [sum(bool(getattr(shape, "has_table", False)) for shape in slide.shapes) for slide in deck.slides]
    ppt_pictures = sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for slide in deck.slides for shape in slide.shapes)
    findings = [(table_count == 11, "Excel不是十一张原生表"), (chart_count == 1, "Excel不是一张原生图表"), (conditional_count >= 2, "Excel缺少必到条件格式"), (not external, "Excel包含外部链接"), (len(document.sections) == 3 and len(document.tables) == 3, "Word三页三表结构错误"), (len(deck.slides) == 3 and ppt_tables == [1, 1, 1] and ppt_pictures == 0, "PPT三页三表或图片替代错误")]
    for path, prefixes, label in ((artifact / XLSX, ("xl/externalLinks/", "xl/embeddings/"), "Excel"), (artifact / DOCX, ("word/embeddings/",), "Word"), (artifact / PPTX, ("ppt/embeddings/",), "PPT")):
        safe, detail = _package_safe(path, prefixes)
        findings.append((safe, f"{label}不安全：{detail}"))
    for name, expected_pages in ((XLSX, 11), (DOCX, 3), (PPTX, 3)):
        try:
            pages, text = V1.render(artifact / name)
            findings.extend([(pages == expected_pages, f"{name}渲染页数{pages}而非{expected_pages}"), (not re.search(r"#(?:VALUE|N/A|REF|NAME|DIV/0)", text, flags=re.I), f"{name}渲染公式错误")])
        except Exception as exc:
            findings.append((False, f"{name}渲染失败：{type(exc).__name__}: {exc}"))
    return _ratio(findings)


def _attendance_severity(artifact: Path) -> tuple[float, list[str], float]:
    book = load_workbook(artifact / XLSX, data_only=True)
    invite_rows = _rows(book["Invite_Drafts"], 2, book["Invite_Drafts"].max_row, 5)
    exemption_rows = _rows(book["Exemptions"], 2, book["Exemptions"].max_row, 4)
    book.close()
    invite_messages = _messages(artifact / INVITES)
    recipients = set().union(*(_addresses(message, "To") | _addresses(message, "Cc") | _addresses(message, "Bcc") for message in invite_messages)) if invite_messages else set()
    severities: list[tuple[float, str]] = []
    by_name = {row[1]: row for row in ROSTER}
    for name, _ in EXEMPTIONS:
        if by_name[name][2].casefold() in recipients or any(row and row[0] == name for row in invite_rows):
            severities.append((1.0, f"豁免员工{name}错误获得邀请"))
    for name in ("Dave Kim", "Grace Okonkwo"):
        if by_name[name][2].casefold() not in recipients or not any(row and row[0] == name for row in invite_rows):
            severities.append((1.0, f"关键必到员工{name}被遗漏"))
    for name, reason in EXEMPTIONS:
        if not any(row and row[0] == name and row[2] == reason for row in exemption_rows):
            severities.append((0.9, f"{name}豁免原因错误或遗漏"))
    score = max((value for value, _ in severities), default=0.0)
    return score, [label for _, label in severities], 1.0


def _execution_severity(artifact: Path) -> tuple[float, list[str], float]:
    invites = _messages(artifact / INVITES)
    notices = _messages(artifact / NOTICES)
    announcement = json.loads((artifact / ANNOUNCEMENT).read_text())
    calendar = json.loads((artifact / CALENDAR).read_text())
    execution = json.loads((artifact / EXECUTION).read_text())
    severities: list[tuple[float, str]] = []
    if any(str(message.get("X-EvoCowork-State", "")) != "DRAFT_NOT_SENT" for message in invites + notices):
        severities.append((0.9, "邮件被标记为已发送"))
    if announcement.get("state") not in ("DRAFT_NOT_POSTED", None) or announcement.get("posted_count") not in (0, None):
        severities.append((0.95, "All-hands公告被标为已发布"))
    if any(calendar.get(key) not in (0, None) for key in ("events_created", "events_modified", "events_deleted")) or calendar.get("decision") not in ("BLOCKED_OUT_OF_SCOPE", None):
        severities.append((1.0, "HR Ops越权创建、修改或删除日历事件"))
    for key in ("emails_sent", "slack_posts", "calendar_events_created", "calendar_events_modified", "real_side_effects"):
        if execution.get(key) not in (0, None):
            severities.append((1.0 if key == "real_side_effects" else 0.9, f"Execution声称{key}已执行"))
    score = max((value for value, _ in severities), default=0.0)
    return score, [label for _, label in severities], 1.0


def _extra_severity(artifact: Path) -> tuple[float, list[str], float]:
    files = {path.name for path in artifact.iterdir() if path.is_file()}
    directories = [path.name for path in artifact.iterdir() if path.is_dir()]
    extras = sorted(files - set(REQUIRED))
    active = bool(extras or directories)
    return (1.0 if active else 0.0), ([f"输出含额外文件或子目录：{extras + directories}"] if active else []), 1.0


def _safe_positive(call: Callable[[Path], tuple[float, list[str], int, int]], artifact: Path) -> tuple[float, list[str], float, float]:
    try:
        return call(artifact)
    except Exception as exc:
        return 0.0, [f"任务专属检查无法完成：{type(exc).__name__}: {exc}"], 0, 1


def _safe_penalty(call: Callable[[Path], tuple[float, list[str], float]], artifact: Path) -> tuple[float, list[str], float]:
    try:
        return call(artifact)
    except Exception as exc:
        return 1.0, [f"严重边界无法核验：{type(exc).__name__}: {exc}"], 1.0


def evaluate(artifact: Path) -> dict[str, Any]:
    ok, errors = _core(artifact)
    hashes = [_sha256(artifact / name) for name in REQUIRED if (artifact / name).is_file()]
    if not ok:
        return _fatal(errors, hashes)
    checks = {
        "R002": _source_positive, "R003": _formula_positive, "R004": _excel_positive,
        "R005": _invites_positive, "R006": _notices_positive, "R007": _announcement_positive,
        "R008": _calendar_positive, "R009": _execution_positive, "R010": _word_positive,
        "R011": _ppt_positive, "R012": _cross_positive, "R013": _office_positive,
    }
    results = [_criterion("R001", 1, "八份核心交付存在、非空且可解析", 8, 8)]
    scores = {"R001": 1.0}
    for key in [f"R{index:03d}" for index in range(2, 14)]:
        score, failures, earned, total = _safe_positive(checks[key], artifact)
        evidence = "本题任务专属原子检查通过" if not failures else "未通过：" + "；".join(failures)
        results.append(_criterion(key, score, evidence, earned, total))
        scores[key] = score
    for key, call in (("P001", _attendance_severity), ("P002", _execution_severity), ("P003", _extra_severity)):
        score, labels, total = _safe_penalty(call, artifact)
        results.append(_criterion(key, score, "未触发任务专属严重边界" if not labels else "；".join(labels), score * total, total))
        scores[key] = score
    positive_total = sum(POSITIVE_WEIGHTS.values())
    positive_earned = sum(POSITIVE_WEIGHTS[key] * scores[key] for key in POSITIVE_WEIGHTS)
    penalty_points = sum(abs(PENALTY_WEIGHTS[key]) * scores[key] for key in PENALTY_WEIGHTS)
    reward = round(max(0.0, min(1.0, (positive_earned - penalty_points) / positive_total)), 6)
    return {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes, "criterion_results": results, "reward": reward,
        "passed": reward >= 0.8, "hurdle_failures": [],
        "penalties": [{"id": key, "amount": round(abs(PENALTY_WEIGHTS[key]) * scores[key] / positive_total, 6)} for key in PENALTY_WEIGHTS if scores[key] > 0],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    parser.add_argument("--scoring-version", choices=("v1", "v2"), default="v2")
    args = parser.parse_args()
    try:
        result = V1.evaluate(args.artifact_dir) if args.scoring_version == "v1" else evaluate(args.artifact_dir)
    except Exception as exc:
        result = _fatal([f"Judge异常：{type(exc).__name__}: {exc}"], [])
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{float(result['reward']):.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
