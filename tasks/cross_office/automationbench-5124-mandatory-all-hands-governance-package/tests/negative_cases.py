#!/usr/bin/env python3
"""AutomationBench 5124 的任务专属 v2 分层负例。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from typing import Callable

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
OUTPUTS = (
    "All_Hands_Attendance_Audit.xlsx", "All_Hands_Operations_Brief.docx", "All_Hands_Operations_Review.pptx",
    "All_Hands_Invitation_Drafts.mbox", "Exempt_Recording_Notice_Drafts.mbox", "All_Hands_Announcement_Draft.json",
    "Calendar_Change_Control.json", "Meeting_Execution_Control.json",
)
XLSX, DOCX, PPTX, INVITES, NOTICES, ANNOUNCEMENT, CALENDAR, EXECUTION = OUTPUTS
Mutator = Callable[[Path], None]


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"无法载入{path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _edit(sheet: str, cell: str, value) -> Mutator:
    def mutate(root: Path) -> None:
        book = load_workbook(root / XLSX)
        book[sheet][cell] = value
        book.save(root / XLSX)
    return mutate


def _json(filename: str, mutator: Callable[[dict], None]) -> Mutator:
    def mutate(root: Path) -> None:
        path = root / filename
        payload = json.loads(path.read_text(encoding="utf-8"))
        mutator(payload)
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return mutate


def _rewrite_message(filename: str, index: int, mutator: Callable) -> Mutator:
    def mutate(root: Path) -> None:
        path = root / filename
        box = mailbox.mbox(path, create=False)
        messages = list(box)
        box.close()
        mutator(messages[index])
        with path.open("wb") as handle:
            for order, message in enumerate(messages):
                handle.write(f"From x Wed Apr 15 09:00:{order:02d} 2026\n".encode())
                BytesGenerator(handle, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
                handle.write(b"\n")
    return mutate


def _replace_body(message, text: str) -> None:
    for header in ("Content-Type", "Content-Transfer-Encoding", "MIME-Version"):
        if header in message:
            del message[header]
    message.set_payload(text)


def _word_dave(root: Path) -> None:
    path = root / DOCX
    document = Document(path)
    for row in document.tables[0].rows:
        if row.cells[0].text == "Dave Kim":
            row.cells[2].text = "EXEMPT_SUPPORT_ON_SHIFT"
            break
    document.save(path)


def _ppt_grace(root: Path) -> None:
    path = root / PPTX
    deck = Presentation(path)
    for row in deck.slides[1].shapes:
        if getattr(row, "has_table", False):
            for table_row in row.table.rows:
                if table_row.cells[0].text == "Grace Okonkwo":
                    table_row.cells[1].text = "EXEMPT"
                    table_row.cells[2].text = "Other obligation"
    deck.save(path)


def _cases() -> list[tuple[str, Mutator, str]]:
    cases: list[tuple[str, Mutator, str]] = []
    for filename in OUTPUTS:
        cases.append((f"missing-{filename}", lambda root, item=filename: (root / item).unlink(), "fatal_zero"))
        cases.append((f"corrupt-{filename}", lambda root, item=filename: (root / item).write_bytes(b"not-a-valid-artifact"), "fatal_zero"))
    cases.extend([
        ("source-dave-note-local", _edit("Source_Roster", "G5", "On shift Wednesday 2-3 PM"), "partial"),
        ("policy-vp-override-local", _edit("Source_Policy", "C2", "Policy text missing VP override"), "partial"),
        ("evidence-roster-hash-local", _edit("Evidence_Register", "C3", "0" * 64), "partial"),
        ("formula-dave-on-shift-local", _edit("Attendance_Decisions", "J5", "YES"), "partial"),
        ("excel-invite-carol-meeting-local", _edit("Invite_Drafts", "D3", "2026-04-23 2-3 PM ET"), "partial"),
        ("excel-elena-reason-local", _edit("Exemptions", "C3", "EXEMPT_SUPPORT_ON_SHIFT"), "partial"),
        ("invite-alice-subject-local", _rewrite_message(INVITES, 0, lambda message: message.replace_header("Subject", "Mandatory all-hands")), "partial"),
        ("invite-bob-forbidden-route", _rewrite_message(INVITES, 0, lambda message: message.replace_header("To", "bob.chen@company.example.com")), "partial"),
        ("notice-elena-reason-local", _rewrite_message(NOTICES, 1, lambda message: _replace_body(message, "Hello Elena Vasquez,\n\nYou are exempt because: EXEMPT_SUPPORT_ON_SHIFT. The meeting recording will be shared after the event. DRAFT_NOT_SENT.\n")), "partial"),
        ("announcement-channel-local", _json(ANNOUNCEMENT, lambda payload: payload.update({"channel_id": "C_WRONG"})), "partial"),
        ("calendar-modified-severe", _json(CALENDAR, lambda payload: payload.update({"events_modified": 1})), "partial"),
        ("execution-required-count-local", _json(EXECUTION, lambda payload: payload.update({"required_live": 4})), "partial"),
        ("execution-email-sent-severe", _json(EXECUTION, lambda payload: payload.update({"emails_sent": 5})), "partial"),
        ("word-dave-decision-local", _word_dave, "partial"),
        ("ppt-grace-override-local", _ppt_grace, "partial"),
        ("extra-file", lambda root: (root / "unexpected.txt").write_text("unexpected", encoding="utf-8"), "partial"),
    ])
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    judge = _load("all_hands_5124_v2_judge", ROOT / "tests" / "judge.py")
    judge.V1.render = lambda path: ({".xlsx": 11, ".docx": 3, ".pptx": 3}[path.suffix.casefold()], "render-ok")
    if judge.evaluate(GOLDEN).get("reward") != 1.0:
        raise AssertionError("Golden v2必须为1.0")
    results = []
    for name, mutator, layer in _cases():
        with tempfile.TemporaryDirectory(prefix=f"all-hands-5124-v2-{name}-") as temporary:
            output = Path(temporary) / "output"
            shutil.copytree(GOLDEN, output)
            mutator(output)
            result = judge.evaluate(output)
            reward = float(result.get("reward", 0))
            if layer == "fatal_zero" and reward != 0:
                raise AssertionError(f"{name}应硬归零，实际{reward}")
            if layer == "partial" and not (0 < reward < 1):
                raise AssertionError(f"{name}应保留连续部分分，实际{reward}")
            results.append({"case": name, "layer": layer, "reward": reward, "passed": bool(result.get("passed")), "penalties": result.get("penalties", [])})
    partial = [item["reward"] for item in results if item["layer"] == "partial"]
    report = {"total": len(results), "fatal_zero": sum(item["layer"] == "fatal_zero" for item in results), "partial_continuous": len(partial), "partial_reward_range": [min(partial), max(partial)], "cases": results}
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
