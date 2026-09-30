#!/usr/bin/env python3
import json,mailbox,shutil,subprocess,sys,tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1];GOLDEN=ROOT/"solution/Golden";PYTHON=sys.executable
def score(case):
 result=case.parent/(case.name+".json");reward=case.parent/(case.name+".txt");subprocess.run([PYTHON,str(ROOT/"tests/judge.py"),"--artifact-dir",str(case),"--result",str(result),"--reward-file",str(reward)],check=True,capture_output=True,text=True,timeout=240);return float(reward.read_text())
def rewrite(path,mutation):
 box=mailbox.mbox(path,create=False);messages=list(box);box.close();mutation(messages)
 with path.open("wb") as handle:
  for index,message in enumerate(messages):handle.write(f"From x Wed Apr 15 09:00:{index:02d} 2026\n".encode());BytesGenerator(handle,policy=SMTP,mangle_from_=True,maxheaderlen=0).flatten(message);handle.write(b"\n")
def replace_body(message,text):
 for header in ("Content-Type","Content-Transfer-Encoding","MIME-Version"):
  if header in message:del message[header]
 message.set_payload(text)
def main():
 with tempfile.TemporaryDirectory(prefix="allhands5124-neg-") as temp:
  root=Path(temp);names=["missing","wrong_required","bob_invited","grace_omitted","wrong_exempt","wrong_date","wrong_recording","wrong_channel","calendar_modified","emails_sent","formula_removed","word_false","ppt_omit"];cases=[]
  for name in names:case=root/name;shutil.copytree(GOLDEN,case);cases.append(case)
  (cases[0]/"All_Hands_Operations_Review.pptx").unlink()
  payload=json.loads((cases[1]/"Meeting_Execution_Control.json").read_text());payload["required_live"]=4;(cases[1]/"Meeting_Execution_Control.json").write_text(json.dumps(payload))
  rewrite(cases[2]/"All_Hands_Invitation_Drafts.mbox",lambda ms:ms[0].replace_header("To","bob.chen@company.example.com"))
  rewrite(cases[3]/"All_Hands_Invitation_Drafts.mbox",lambda ms:ms.pop())
  rewrite(cases[4]/"Exempt_Recording_Notice_Drafts.mbox",lambda ms:ms[0].replace_header("To","dave.kim@company.example.com"))
  rewrite(cases[5]/"All_Hands_Invitation_Drafts.mbox",lambda ms:replace_body(ms[0],"Meeting is April 23, 4-5 PM ET. DRAFT_NOT_SENT."))
  rewrite(cases[6]/"Exempt_Recording_Notice_Drafts.mbox",lambda ms:ms[0].replace_header("To","grace.okonkwo@company.example.com"))
  payload=json.loads((cases[7]/"All_Hands_Announcement_Draft.json").read_text());payload["channel_id"]="C_WRONG";(cases[7]/"All_Hands_Announcement_Draft.json").write_text(json.dumps(payload))
  payload=json.loads((cases[8]/"Calendar_Change_Control.json").read_text());payload["events_modified"]=2;(cases[8]/"Calendar_Change_Control.json").write_text(json.dumps(payload))
  payload=json.loads((cases[9]/"Meeting_Execution_Control.json").read_text());payload["emails_sent"]=8;(cases[9]/"Meeting_Execution_Control.json").write_text(json.dumps(payload))
  workbook=load_workbook(cases[10]/"All_Hands_Attendance_Audit.xlsx");workbook["Attendance_Decisions"]["M2"]="INVITATION_DRAFT";workbook.save(cases[10]/"All_Hands_Attendance_Audit.xlsx")
  document=Document(cases[11]/"All_Hands_Operations_Brief.docx");document.add_paragraph("Calendar events were rescheduled by HR Ops.");document.save(cases[11]/"All_Hands_Operations_Brief.docx")
  presentation=Presentation(cases[12]/"All_Hands_Operations_Review.pptx")
  for slide in presentation.slides:
   for shape in slide.shapes:
    if getattr(shape,"has_text_frame",False):shape.text=shape.text.replace("Grace Okonkwo","Unknown")
    if getattr(shape,"has_table",False):
     for row in shape.table.rows:
      for cell in row.cells:cell.text=cell.text.replace("Grace Okonkwo","Unknown")
  presentation.save(cases[12]/"All_Hands_Operations_Review.pptx")
  rewards={case.name:score(case) for case in cases};print(json.dumps(rewards,indent=2));assert all(value==0 for value in rewards.values()),rewards
if __name__=="__main__":main()
