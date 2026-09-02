#!/usr/bin/env python3
"""只静态化 AutomationBench 5125；不是通用造题器。"""
import hashlib,importlib.util,json,shutil,sys,types
from datetime import datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.worksheet.table import Table,TableStyleInfo
ROOT=Path(__file__).resolve().parents[1];UP=Path(__file__).with_name("upstream-hr-tasks.py");SNAP=Path(__file__).with_name("upstream-task-5125.json");DATA=ROOT/"data/input_files";ENV=ROOT/"environment/input_files";UPSTREAM_SHA="c173a36a03f43096bd122577584aa32f9df379d8760aa46c43be04c85d101779"
def load_task():
 if hashlib.sha256(UP.read_bytes()).hexdigest()!=UPSTREAM_SHA:raise ValueError("上游哈希错误")
 d=types.ModuleType("datasets");d.Dataset=object;sys.modules["datasets"]=d
 for name in ("automationbench","automationbench.domains","automationbench.domains.hr"):sys.modules[name]=types.ModuleType(name)
 noise=types.ModuleType("automationbench.domains.hr._noise");noise.apply_noise=lambda x:x;sys.modules[noise.__name__]=noise
 spec=importlib.util.spec_from_file_location("automationbench_5125",UP);module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module);task=module.get_task_5125();assert task["example_id"]==5125 and task["task"]=="hr.data_migration_validation";return task
def audit(task):
 info=task["info"];state=info["initial_state"];assert state["meta"]["current_time"]=="2026-04-15T09:00:00Z" and len(info["assertions"])==7
 rows=state["google_sheets"]["spreadsheets"][0]["worksheets"][0]["rows"];assert len(rows)==8 and [r["cells"]["Employee ID"] for r in rows]==["EMP-6601","EMP-6618","EMP-6625","EMP-6632","EMP-6655","EMP-6640","EMP-6648","EMP-6662"]
 assert rows[0]["cells"]["Employee"]==rows[4]["cells"]["Employee"]=="Michael Park" and rows[0]["cells"]["Department"]=="Engineering" and rows[4]["cells"]["Department"]=="Sales"
 body=state["gmail"]["messages"][0]["body_plain"];assert "EMP-6601" in body and "Senior Engineer" in body and "Sarah Kim" in body and "EMP-6655" in body and "record is correct" in body
def write_mbox(messages,path):
 chunks=[]
 for record in messages:
  msg=EmailMessage(policy=SMTP);msg["From"]=record["from_"];msg["To"]=", ".join(record["to"]);msg["Subject"]=record["subject"];dt=datetime.fromisoformat(record["date"].replace("Z","+00:00"));msg["Date"]=dt.strftime("%a, %d %b %Y %H:%M:%S +0000");msg["Message-ID"]=f"<{record['id']}@automationbench.invalid>";msg["X-AutomationBench-Message-ID"]=record["id"];msg.set_content(record["body_plain"],charset="utf-8");buf=BytesIO();BytesGenerator(buf,policy=SMTP,mangle_from_=True,maxheaderlen=0).flatten(msg);chunks.extend([f"From {record['from_']} {dt.strftime('%a %b %d %H:%M:%S %Y')}\n".encode(),buf.getvalue().replace(b"\r\n",b"\n"),b"\n"])
 path.write_bytes(b"".join(chunks))
def write_xlsx(rows,path):
 wb=Workbook();ws=wb.active;ws.title="Migrated Records";headers=["Source Row","Employee","Employee ID","Department","Title","Manager","Validation Status","Notes"];ws.append(headers)
 for row in rows:ws.append([row["row_id"]]+[row["cells"][k] for k in headers[1:]])
 for i,w in enumerate([10,24,18,22,28,24,24,46],1):ws.column_dimensions[ws.cell(1,i).column_letter].width=w
 ws.freeze_panes="A2";ws.sheet_view.showGridLines=False
 for c in ws[1]:c.fill=PatternFill("solid",fgColor="17365D");c.font=Font(name="Carlito",bold=True,color="FFFFFF");c.alignment=Alignment(horizontal="center",wrap_text=True)
 for row in ws.iter_rows(min_row=2):
  for c in row:c.font=Font(name="Carlito",size=9);c.alignment=Alignment(vertical="top",wrap_text=True)
 table=Table(displayName="MigrationTrackerSource",ref=ws.dimensions);table.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True);ws.add_table(table);ws.page_setup.orientation="landscape";ws.page_setup.fitToWidth=1;ws.sheet_properties.pageSetUpPr.fitToPage=True;wb.save(path)
def main():
 task=load_task();audit(task);state=task["info"]["initial_state"];DATA.mkdir(parents=True,exist_ok=True);ENV.mkdir(parents=True,exist_ok=True);snapshot=json.dumps(task,ensure_ascii=False,indent=2)+"\n"
 if SNAP.exists() and SNAP.read_text()!=snapshot:raise ValueError("单题快照漂移")
 SNAP.write_text(snapshot);write_xlsx(state["google_sheets"]["spreadsheets"][0]["worksheets"][0]["rows"],DATA/"Employee_Data_Migration_Tracker.xlsx");write_mbox(state["gmail"]["messages"],DATA/"Post_Migration_Corrections.mbox");(DATA/"HR_Slack_Context.json").write_text(json.dumps(state["slack"],ensure_ascii=False,indent=2)+"\n");(DATA/"Workflow_Context.json").write_text(json.dumps({"meta":state["meta"]},indent=2)+"\n")
 for item in DATA.iterdir():
  if item.is_file():shutil.copy2(item,ENV/item.name)
if __name__=="__main__":main()
