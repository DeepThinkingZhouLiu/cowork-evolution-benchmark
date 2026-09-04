#!/usr/bin/env python3
"""AutomationBench 5125 Employee ID 消歧与迁移差异专属 Judge。"""
import argparse,json,mailbox,re
from pathlib import Path
from openpyxl import load_workbook
from office_checks import body_text,docx_text,package_safe,pptx_text,render,sha256
TASK="cowork-evo/automationbench-5125-data-migration-validation-package";HASHES={"Employee_Data_Migration_Tracker.xlsx":"8820dbcdf0415a15142424fd7d20552aef29bbe8ef8deec4598f623716eacf50","HR_Slack_Context.json":"bd2d19e20559df9ee06a07f8164d5dc7fce2ab417c6e0f7cc982fc2630aea7c8","Post_Migration_Corrections.mbox":"358991e170a8eb1707faf83bebb90662895f0a86291168911e5c13cdd5b8f738","Workflow_Context.json":"76c4a0657bc27f2fd830b10613ec7ceae3adcdc5798aeee0f822d27bf69bbc99"}
X="Employee_Data_Migration_Audit.xlsx";D="Data_Migration_Validation_Brief.docx";P="Data_Migration_Validation_Review.pptx";M="HRIS_Correction_Draft.mbox";U="Validation_Status_Update_Drafts.json";G="Migration_Discrepancy_Register.json";C="Migration_Execution_Control.json";REQUIRED={X,D,P,M,U,G,C};WEIGHTS={f"R{i:03d}":v for i,v in enumerate([6,8,10,10,10,8,9,8,8,8,7,8],1)};CAPS={k:(0 if k=="R001" else .25) for k in WEIGHTS if k!="R012"}
class Checks:
 def __init__(self):self.v={k:{"passed":True,"details":[]} for k in WEIGHTS}
 def require(self,key,condition,detail):
  if not condition:self.v[key]["passed"]=False;self.v[key]["details"].append(detail) if detail not in self.v[key]["details"] else None
 def ok(self,key):return self.v[key]["passed"]
def root():return Path(__file__).resolve().parents[1]
def inp():return Path("/app/input_files") if Path("/app/input_files").is_dir() else root()/"data/input_files"
def rows(ws):return [tuple(c.value for c in row) for row in ws.iter_rows(min_row=2)]
def source(q):
 for name,digest in HASHES.items():q.require("R001",(inp()/name).is_file() and sha256(inp()/name)==digest,f"输入{name}")
 wb=load_workbook(inp()/"Employee_Data_Migration_Tracker.xlsx",read_only=True,data_only=True);r=rows(wb["Migrated Records"]);q.require("R002",[x[2] for x in r]==["EMP-6601","EMP-6618","EMP-6625","EMP-6632","EMP-6655","EMP-6640","EMP-6648","EMP-6662"],"8个源ID")
 q.require("R002",r[0][1]==r[4][1]=="Michael Park" and r[0][3:5]==("Engineering","Engineer") and r[4][3:5]==("Sales","Sales Rep"),"同名Michael陷阱")
 msg=list(mailbox.mbox(inp()/"Post_Migration_Corrections.mbox",create=False));text=body_text(msg[0]) if len(msg)==1 else "";q.require("R002",all(t in text for t in ("EMP-6601","Senior Engineer","Sarah Kim","Product","EMP-6655","record is correct")),"更正邮件事实")
def workbook(path,q):
 wb=load_workbook(path,data_only=False);expected=["Source_Tracker","Source_Corrections","Evidence_Register","Validation_Decisions","Discrepancies","Validated_Records","Status_Update_Drafts","HRIS_Email_Draft","Reconciliation","Control_Summary"];q.require("R003",wb.sheetnames==expected,"10表")
 formulas=[c.value for row in wb["Validation_Decisions"].iter_rows(min_row=2) for c in row if isinstance(c.value,str) and c.value.startswith("=")];q.require("R004",len(formulas)==104,f"公式{len(formulas)}");joined="\n".join(formulas)
 for token in ('C2="EMP-6618"','C2="EMP-6601"','Senior Engineer','Product','DISCREPANCY','VALIDATED','Source_Tracker!H9'):q.require("R004",token in joined,f"公式缺{token}")
 discrepancies=rows(wb["Discrepancies"]);q.require("R005",discrepancies==[("Michael Park","EMP-6601","Title","Engineer","Senior Engineer","Promoted in March"),("Sarah Kim","EMP-6618","Department","Engineering","Product","Transferred April 1")],"2条差异")
 valid=rows(wb["Validated_Records"]);q.require("R005",len(valid)==6 and [r[1] for r in valid]==["EMP-6625","EMP-6632","EMP-6655","EMP-6640","EMP-6648","EMP-6662"] and next(r for r in valid if r[1]=="EMP-6655")[2:4]==("Sales","Sales Rep"),"6条验证及EMP-6655")
 recon=rows(wb["Reconciliation"]);q.require("R006",("Duplicate-name records",2,"Disambiguated by Employee ID") in recon and ("Source rows updated",0,"Draft only") in recon,"ID消歧与零更新")
def artifacts(artifact,q):
 msgs=list(mailbox.mbox(artifact/M,create=False));text=body_text(msgs[0]) if len(msgs)==1 else "";q.require("R007",len(msgs)==1 and str(msgs[0].get("To"))=="hris@company.example.com" and all(t in text for t in ("EMP-6601","Senior Engineer","Sarah Kim","EMP-6618","Product","DRAFT_NOT_SENT")) and "EMP-6655" not in text,"更正草稿仅两条")
 updates=json.loads((artifact/U).read_text());items=updates.get("updates",[]);q.require("R008",updates.get("draft_count")==8 and updates.get("applied_count")==0 and updates.get("state")=="DRAFT_NOT_APPLIED" and [(x["employee_id"],x["validation_status"],x["applied"]) for x in items]==[("EMP-6601","Discrepancy",False),("EMP-6618","Discrepancy",False),("EMP-6625","Validated",False),("EMP-6632","Validated",False),("EMP-6655","Validated",False),("EMP-6640","Validated",False),("EMP-6648","Validated",False),("EMP-6662","Validated",False)],"8条状态更新草稿")
 register=json.loads((artifact/G).read_text());q.require("R008",register.get("discrepancy_count")==2 and register.get("validated_count")==6 and [(x["employee_id"],x["field"],x["expected"]) for x in register.get("discrepancies",[])]==[("EMP-6601","Title","Senior Engineer"),("EMP-6618","Department","Product")],"差异登记")
 control=json.loads((artifact/C).read_text());expected={"records_reviewed":8,"discrepancies":2,"validated_records":6,"status_update_drafts":8,"status_updates_applied":0,"correction_email_drafts":1,"emails_sent":0,"source_rows_updated":0,"real_side_effects":0};q.require("R009",all(control.get(k)==v for k,v in expected.items()),"执行边界")
def office(artifact,q):
 text,sections,tables=docx_text(artifact/D);q.require("R010",sections==3 and tables==3,"Word结构")
 for token in ("Eight migrated","EMP-6601","Senior Engineer","Sarah Kim","EMP-6618","Product","EMP-6655","Sales Rep","DRAFT_NOT_APPLIED"):q.require("R010",token.lower() in text.lower(),f"Word缺{token}")
 q.require("R010","emp-6655 requires correction" not in text.lower(),"Word不得误报EMP-6655")
 text,slides,tables,pictures=pptx_text(artifact/P);q.require("R011",slides==3 and tables==[1,1,1] and pictures==0,"PPT结构")
 for token in ("2 discrepancies","6 validated","EMP-6601","EMP-6618","EMP-6655","Employee ID prevents","DRAFT_NOT_APPLIED"):q.require("R011",token in text,f"PPT缺{token}")
 for name in (X,D,P):safe,detail=package_safe(artifact/name);q.require("R012",safe,f"安全{name}:{detail}")
def finish(q,pages):
 hurdles=[k for k in CAPS if not q.ok(k)];base=sum(v for k,v in WEIGHTS.items() if q.ok(k))/100;cap=min([1]+[CAPS[k] for k in hurdles]);penalty=(.3 if any(not q.ok(k) for k in ("R002","R003","R004","R005","R006")) else 0)+(.3 if any(not q.ok(k) for k in ("R007","R008","R009")) else 0)+(.25 if any(not q.ok(k) for k in ("R001","R010","R011","R012")) else 0);reward=round(max(0,min(base,cap)-penalty),6);return {"task_id":TASK,"reward":reward,"passed":reward>=.8 and not hurdles,"criteria":q.v,"hurdle_failures":hurdles,"source_truth":{"reviewed":8,"discrepancies":2,"validated":6,"source_updates":0,"real_side_effects":0},"render_pages":pages}
def evaluate(artifact):
 q=Checks();source(q);files={p.name for p in artifact.iterdir() if p.is_file()} if artifact.is_dir() else set();q.require("R001",files==REQUIRED,f"集合{files}")
 if files!=REQUIRED:return finish(q,{})
 for fn,keys in ((lambda:workbook(artifact/X,q),("R003","R004","R005","R006","R012")),(lambda:artifacts(artifact,q),("R007","R008","R009")),(lambda:office(artifact,q),("R010","R011","R012"))):
  try:fn()
  except Exception as exc:
   for key in keys:q.require(key,False,str(exc))
 pages={}
 for name,expected in ((X,10),(D,3),(P,3)):
  try:count,text=render(artifact/name);pages[name]=count;q.require("R012",count==expected,f"页数{name}:{count}");q.require("R012",not re.search(r"#(?:VALUE|REF|NAME|DIV/0|N/A)",text,re.I),f"公式错{name}")
  except Exception as exc:q.require("R012",False,str(exc))
 return finish(q,pages)
def main():
 p=argparse.ArgumentParser();p.add_argument("--artifact-dir",type=Path,required=True);p.add_argument("--result",type=Path,required=True);p.add_argument("--reward-file",type=Path,required=True);a=p.parse_args();r=evaluate(a.artifact_dir);a.result.parent.mkdir(parents=True,exist_ok=True);a.reward_file.parent.mkdir(parents=True,exist_ok=True);a.result.write_text(json.dumps(r,ensure_ascii=False,indent=2)+"\n");a.reward_file.write_text(f"{r['reward']:.6f}\n");print(json.dumps(r,ensure_ascii=False,indent=2))
if __name__=="__main__":main()
