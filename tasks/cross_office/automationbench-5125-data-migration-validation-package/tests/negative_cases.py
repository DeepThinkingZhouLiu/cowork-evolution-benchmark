#!/usr/bin/env python3
"""AutomationBench 5125 的任务专属 v2 分层负例。"""
from __future__ import annotations
import argparse,importlib.util,json,mailbox,shutil,sys,tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from typing import Callable
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1];GOLDEN=ROOT/"solution/Golden";INPUTS=ROOT/"environment/input_files"
OUTPUTS=("Employee_Data_Migration_Audit.xlsx","Data_Migration_Validation_Brief.docx","Data_Migration_Validation_Review.pptx","HRIS_Correction_Draft.mbox","Validation_Status_Update_Drafts.json","Migration_Discrepancy_Register.json","Migration_Execution_Control.json")
XLSX,DOCX,PPTX,MBOX,UPDATES,REGISTER,EXECUTION=OUTPUTS;Mutator=Callable[[Path],None]
def _load(name,path):
 spec=importlib.util.spec_from_file_location(name,path)
 if spec is None or spec.loader is None:raise RuntimeError(f"无法载入{path}")
 module=importlib.util.module_from_spec(spec);sys.modules[spec.name]=module;spec.loader.exec_module(module);return module
def _input_only(root):
 for path in list(root.iterdir()):
  if path.is_file():path.unlink()
 for source in INPUTS.iterdir():
  if source.is_file():shutil.copy2(source,root/source.name)
def _edit(sheet,cell,value):
 def mutate(root):
  book=load_workbook(root/XLSX);book[sheet][cell]=value;book.save(root/XLSX)
 return mutate
def _json(filename,mutator):
 def mutate(root):
  path=root/filename;payload=json.loads(path.read_text());mutator(payload);path.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n")
 return mutate
def _message(mutator):
 def mutate(root):
  path=root/MBOX;box=mailbox.mbox(path,create=False);messages=list(box);box.close();mutator(messages[0])
  with path.open("wb") as handle:
   handle.write(b"From x Wed Apr 15 09:00:00 2026\n");BytesGenerator(handle,policy=SMTP,mangle_from_=True,maxheaderlen=0).flatten(messages[0]);handle.write(b"\n")
 return mutate
def _replace_body(message,text):
 for header in ("Content-Type","Content-Transfer-Encoding","MIME-Version"):
  if header in message:del message[header]
 message.set_payload(text)
def _word_false(root):
 path=root/DOCX;document=Document(path);document.add_paragraph("EMP-6655 requires correction.");document.save(path)
def _ppt_emp6655(root):
 path=root/PPTX;deck=Presentation(path)
 for slide in deck.slides:
  for shape in slide.shapes:
   if getattr(shape,"has_text_frame",False):shape.text=shape.text.replace("EMP-6655","EMP-6601")
   if getattr(shape,"has_table",False):
    for row in shape.table.rows:
     for cell in row.cells:cell.text=cell.text.replace("EMP-6655","EMP-6601")
 deck.save(path)
def _cases():
 cases=[]
 for filename in OUTPUTS:
  cases.append((f"missing-{filename}",lambda root,item=filename:(root/item).unlink(),"fatal_zero"));cases.append((f"corrupt-{filename}",lambda root,item=filename:(root/item).write_bytes(b"not-a-valid-artifact"),"fatal_zero"))
 cases.extend([
  ("mbox-empty",lambda root:(root/MBOX).write_bytes(b""),"fatal_zero"),("unmodified-input-only",_input_only,"fatal_zero"),
  ("source-emp6655-id-local",_edit("Source_Tracker","C6","EMP-6601"),"partial"),("correction-source-local",_edit("Source_Corrections","B2","Corrections"),"partial"),
  ("evidence-tracker-hash-local",_edit("Evidence_Register","C2","0"*64),"partial"),("formula-emp6601-title-local",_edit("Validation_Decisions","J2","Engineer"),"partial"),
  ("excel-discrepancy-title-local",_edit("Discrepancies","E2","Principal Engineer"),"partial"),("excel-emp6655-false-discrepancy",_edit("Validated_Records","D4","Senior Engineer"),"partial"),
  ("mbox-subject-local",_message(lambda message:message.replace_header("Subject","Migration corrections")),"partial"),
  ("mbox-emp6655-forbidden",_message(lambda message:_replace_body(message,"EMP-6601 -> Senior Engineer; EMP-6618 -> Product; EMP-6655 requires correction. DRAFT_NOT_SENT.\n")),"partial"),
  ("updates-emp6655-status-local",_json(UPDATES,lambda payload:payload["updates"][4].update({"validation_status":"Discrepancy"})),"partial"),
  ("updates-applied-severe",_json(UPDATES,lambda payload:payload.update({"applied_count":8,"state":"APPLIED"})),"partial"),
  ("register-sarah-current-local",_json(REGISTER,lambda payload:payload["discrepancies"][1].update({"current":"Sales"})),"partial"),
  ("execution-count-local",_json(EXECUTION,lambda payload:payload.update({"validated_records":5})),"partial"),
  ("execution-source-updated-severe",_json(EXECUTION,lambda payload:payload.update({"source_rows_updated":2})),"partial"),
  ("word-emp6655-false",_word_false,"partial"),("ppt-emp6655-id-local",_ppt_emp6655,"partial"),("extra-file",lambda root:(root/"unexpected.txt").write_text("unexpected"),"partial")])
 return cases
def main():
 parser=argparse.ArgumentParser();parser.add_argument("--report",type=Path);args=parser.parse_args();judge=_load("migration_5125_v2_judge",ROOT/"tests/judge.py");judge.V1.render=lambda path:({".xlsx":10,".docx":3,".pptx":3}[path.suffix.casefold()],"render-ok")
 if judge.evaluate(GOLDEN).get("reward")!=1.0:raise AssertionError("Golden v2必须为1.0")
 results=[]
 for name,mutator,layer in _cases():
  with tempfile.TemporaryDirectory(prefix=f"migration-5125-v2-{name}-") as temporary:
   output=Path(temporary)/"output";shutil.copytree(GOLDEN,output);mutator(output);result=judge.evaluate(output);reward=float(result.get("reward",0))
   if layer=="fatal_zero" and reward!=0:raise AssertionError(f"{name}应硬归零，实际{reward}")
   if layer=="partial" and not 0<reward<1:raise AssertionError(f"{name}应保留连续部分分，实际{reward}")
   results.append({"case":name,"layer":layer,"reward":reward,"passed":bool(result.get("passed")),"penalties":result.get("penalties",[])})
 partial=[item["reward"] for item in results if item["layer"]=="partial"];report={"total":len(results),"fatal_zero":sum(item["layer"]=="fatal_zero" for item in results),"partial_continuous":len(partial),"partial_reward_range":[min(partial),max(partial)],"cases":results}
 if args.report:args.report.parent.mkdir(parents=True,exist_ok=True);args.report.write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n")
 print(json.dumps(report,ensure_ascii=False,indent=2))
if __name__=="__main__":main()
