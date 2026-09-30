#!/usr/bin/env python3
import json,mailbox,shutil,subprocess,sys,tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1];GOLDEN=ROOT/"solution/Golden";PYTHON=sys.executable
def score(case):
 result=case.parent/(case.name+".json");reward=case.parent/(case.name+".txt");subprocess.run([PYTHON,str(ROOT/"tests/judge.py"),"--artifact-dir",str(case),"--result",str(result),"--reward-file",str(reward)],check=True,capture_output=True,text=True,timeout=240);return float(reward.read_text())
def rewrite(path,mutation):
 box=mailbox.mbox(path,create=False);messages=list(box);box.close();mutation(messages)
 with path.open("wb") as handle:
  for i,m in enumerate(messages):handle.write(f"From x Wed Apr 15 09:00:{i:02d} 2026\n".encode());BytesGenerator(handle,policy=SMTP,mangle_from_=True,maxheaderlen=0).flatten(m);handle.write(b"\n")
def replace_body(m,text):
 for h in ("Content-Type","Content-Transfer-Encoding","MIME-Version"):
  if h in m:del m[h]
 m.set_payload(text)
def main():
 with tempfile.TemporaryDirectory(prefix="migration5125-neg-") as temp:
  root=Path(temp);names=["missing","wrong_count","emp6655_corrected","sarah_omitted","wrong_recipient","wrong_status","updates_applied","wrong_register","source_updated","formula_removed","word_false","ppt_omit"];cases=[]
  for n in names:c=root/n;shutil.copytree(GOLDEN,c);cases.append(c)
  (cases[0]/"Data_Migration_Validation_Review.pptx").unlink()
  p=json.loads((cases[1]/"Migration_Execution_Control.json").read_text());p["discrepancies"]=3;(cases[1]/"Migration_Execution_Control.json").write_text(json.dumps(p))
  rewrite(cases[2]/"HRIS_Correction_Draft.mbox",lambda ms:replace_body(ms[0],"EMP-6601 Senior Engineer; EMP-6618 Product; EMP-6655 requires correction. DRAFT_NOT_SENT."))
  rewrite(cases[3]/"HRIS_Correction_Draft.mbox",lambda ms:replace_body(ms[0],"Michael Park EMP-6601 -> Senior Engineer. DRAFT_NOT_SENT."))
  rewrite(cases[4]/"HRIS_Correction_Draft.mbox",lambda ms:ms[0].replace_header("To","wrong@company.example.com"))
  p=json.loads((cases[5]/"Validation_Status_Update_Drafts.json").read_text());next(x for x in p["updates"] if x["employee_id"]=="EMP-6655")["validation_status"]="Discrepancy";(cases[5]/"Validation_Status_Update_Drafts.json").write_text(json.dumps(p))
  p=json.loads((cases[6]/"Validation_Status_Update_Drafts.json").read_text());p["applied_count"]=8;(cases[6]/"Validation_Status_Update_Drafts.json").write_text(json.dumps(p))
  p=json.loads((cases[7]/"Migration_Discrepancy_Register.json").read_text());p["discrepancies"][0]["expected"]="Principal Engineer";(cases[7]/"Migration_Discrepancy_Register.json").write_text(json.dumps(p))
  p=json.loads((cases[8]/"Migration_Execution_Control.json").read_text());p["source_rows_updated"]=8;(cases[8]/"Migration_Execution_Control.json").write_text(json.dumps(p))
  w=load_workbook(cases[9]/"Employee_Data_Migration_Audit.xlsx");w["Validation_Decisions"]["K2"]="DISCREPANCY";w.save(cases[9]/"Employee_Data_Migration_Audit.xlsx")
  d=Document(cases[10]/"Data_Migration_Validation_Brief.docx");d.add_paragraph("EMP-6655 requires correction.");d.save(cases[10]/"Data_Migration_Validation_Brief.docx")
  p=Presentation(cases[11]/"Data_Migration_Validation_Review.pptx")
  for s in p.slides:
   for shape in s.shapes:
    if getattr(shape,"has_text_frame",False):shape.text=shape.text.replace("EMP-6655","UNKNOWN")
    if getattr(shape,"has_table",False):
     for row in shape.table.rows:
      for cell in row.cells:cell.text=cell.text.replace("EMP-6655","UNKNOWN")
  p.save(cases[11]/"Data_Migration_Validation_Review.pptx")
  rewards={c.name:score(c) for c in cases};print(json.dumps(rewards,indent=2));assert all(v==0 for v in rewards.values()),rewards
if __name__=="__main__":main()
