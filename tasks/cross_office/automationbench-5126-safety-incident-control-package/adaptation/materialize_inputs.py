#!/usr/bin/env python3
"""只把 AutomationBench 5126 的公开初始状态静态化为本题输入。

脚本逐项断言八条安全事故、七章安全手册、冻结时间、
Safety Slack 快照和八条上游 assertions。它不是通用造题器。
"""

from __future__ import annotations

import hashlib
import importlib.util
import json
import shutil
import sys
import types
import zipfile
from datetime import date, datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
UPSTREAM = Path(__file__).with_name("upstream-hr-tasks.py")
SNAPSHOT = Path(__file__).with_name("upstream-task-5126.json")
LICENSE = Path(__file__).with_name("LICENSE-AutomationBench-MIT")
DATA_DIR = ROOT / "data" / "input_files"
ENV_DIR = ROOT / "environment" / "input_files"
FIXED_ZIP_TIME = (2026, 4, 15, 9, 0, 0)
SOURCE_SHA256 = "c173a36a03f43096bd122577584aa32f9df379d8760aa46c43be04c85d101779"
LICENSE_SHA256 = "c6f9d87aff6d627d5cd78cb09364dafed7d6f1c4bb7a91183d57377157e96ba4"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_task() -> dict:
    if sha256(UPSTREAM) != SOURCE_SHA256 or sha256(LICENSE) != LICENSE_SHA256:
        raise ValueError("冻结的 AutomationBench 源文件或 LICENSE 哈希错误")
    datasets = types.ModuleType("datasets")
    datasets.Dataset = object
    sys.modules["datasets"] = datasets
    for name in ("automationbench", "automationbench.domains", "automationbench.domains.hr"):
        sys.modules[name] = types.ModuleType(name)
    noise = types.ModuleType("automationbench.domains.hr._noise")
    noise.apply_noise = lambda value: value
    sys.modules[noise.__name__] = noise
    spec = importlib.util.spec_from_file_location("automationbench_hr_tasks_frozen_5126", UPSTREAM)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入冻结的 AutomationBench HR 源文件")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    task = module.get_task_5126()
    if task.get("example_id") != 5126 or task.get("task") != "hr.safety_incident_reporting":
        raise ValueError("冻结函数不再是 AutomationBench 5126")
    return task


def assert_shape(task: dict) -> None:
    prompt = task["prompt"][1]["content"]
    assert all(token in prompt for token in (
        "workplace safety incident reports", "Route and document each one",
        "Google Sheets and Gmail", "#safety on Slack",
    ))
    state = task["info"]["initial_state"]
    assert state["meta"] == {"schema_version":"0.1.0","current_time":"2026-04-15T09:00:00Z"}
    messages = state["gmail"]["messages"]
    assert [item["id"] for item in messages] == ["msg_safety_manual_5126"]
    manual = messages[0]["body_plain"]
    assert all(token in manual for token in (
        "Minor: First aid only", "Moderate: Medical treatment needed",
        "Serious: Hospitalization", "Minor: Email the employee's manager",
        "Moderate: Email the Safety Officer", "Serious: Email the Safety Officer",
        "workers' comp claim form", "reported to OSHA within 24 hours",
        "hospitalization, amputation, or loss of an eye", "legal@company.example.com",
        "retained for 5 years",
    ))
    spreadsheet = state["google_sheets"]["spreadsheets"][0]
    assert spreadsheet["id"] == "ss_safety_5126" and spreadsheet["title"] == "Safety Incident Tracker"
    worksheet = spreadsheet["worksheets"][0]
    assert worksheet["id"] == "ws_incidents_5126" and worksheet["title"] == "Recent Incidents"
    rows = worksheet["rows"]
    assert [item["row_id"] for item in rows] == list(range(2, 10))
    assert [item["cells"]["Employee"] for item in rows] == [
        "Ana Petrova", "Ben Torres", "Carlos Mendoza", "Diana Osei",
        "Ethan Park", "Fiona Strand", "George Tan", "Holly Ng",
    ]
    assert [item["cells"]["Severity"] for item in rows] == ["Minor","Moderate","Serious","Minor","Serious","Minor","Moderate","Serious"]
    assert [item["cells"]["Lost Days"] for item in rows] == ["0","2","TBD","0","5+","0","1","TBD"]
    assert [item["cells"]["Status"] for item in rows] == ["Pending"] * 8
    assert "Hospitalization - emergency surgery" == rows[2]["cells"]["Treatment"]
    assert "Hospitalization - potential loss of vision" == rows[4]["cells"]["Treatment"]
    assert "Hospitalization - surgical amputation" == rows[7]["cells"]["Treatment"]
    assert state["slack"] == {"channels":[{"id":"C_SAF_5126","name":"safety","is_private":False}],"messages":[]}
    assertions = task["info"]["assertions"]
    assert len(assertions) == 8
    serialized = json.dumps(assertions)
    assert all(token in serialized for token in (
        "Ben Torres", "Carlos Mendoza", "Ethan Park",
        "Holly Ng", "George Tan", "Diana Osei", "legal@company.example.com",
    ))


def normalize_xlsx(path: Path) -> None:
    temporary = path.with_suffix(".normalized.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    temporary.replace(path)


def style_sheet(sheet, widths: dict[str, float]) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = sheet.dimensions


def write_tracker(state: dict, target: Path) -> None:
    source = state["google_sheets"]["spreadsheets"][0]
    worksheet = source["worksheets"][0]
    workbook = Workbook()
    workbook.creator = "AutomationBench 5126 public fixture"
    workbook.title = source["title"]
    workbook.subject = "Static Office representation of safety incident tracker"
    workbook.description = "No employee, incident, treatment, severity, date, manager or status fact was added beyond AutomationBench 5126."
    sheet = workbook.active
    sheet.title = worksheet["title"]
    headers = ["Source Row","Employee","Manager Email","Incident Date","Description","Severity","Treatment","Lost Days","Status","Notes"]
    sheet.append(headers)
    for item in worksheet["rows"]:
        cells = item["cells"]
        sheet.append([
            item["row_id"], cells["Employee"], cells["Manager Email"],
            date.fromisoformat(cells["Incident Date"]), cells["Description"], cells["Severity"],
            cells["Treatment"], cells["Lost Days"], cells["Status"], cells["Notes"],
        ])
    for row in range(2, 10):
        sheet.cell(row, 4).number_format = "yyyy-mm-dd"
        sheet.row_dimensions[row].height = 42
    table = Table(displayName="SafetyIncidentSourceInput", ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    style_sheet(sheet, {"A":12,"B":22,"C":30,"D":18,"E":42,"F":14,"G":44,"H":13,"I":14,"J":20})
    workbook.save(target)
    normalize_xlsx(target)


def message_bytes(record: dict) -> bytes:
    message = EmailMessage(policy=SMTP)
    message["From"] = record["from_"]
    message["To"] = ", ".join(record["to"])
    message["Subject"] = record["subject"]
    parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
    message["Date"] = parsed.strftime("%a, %d %b %Y %H:%M:%S +0000")
    message["Message-ID"] = f"<{record['id']}@automationbench.invalid>"
    message["X-AutomationBench-Message-ID"] = record["id"]
    message["X-AutomationBench-Thread-ID"] = record["thread_id"]
    message["X-AutomationBench-Labels"] = ",".join(record.get("label_ids", []))
    message.set_content(record["body_plain"], charset="utf-8")
    buffer = BytesIO()
    BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
    return buffer.getvalue()


def write_mbox(messages: list[dict], target: Path) -> None:
    chunks: list[bytes] = []
    for record in messages:
        parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
        chunks.append(f"From {record['from_']} {parsed.strftime('%a %b %d %H:%M:%S %Y')}\n".encode("ascii"))
        chunks.append(message_bytes(record).replace(b"\r\n", b"\n"))
        chunks.append(b"\n")
    target.write_bytes(b"".join(chunks))


def main() -> None:
    task = load_task()
    assert_shape(task)
    state = task["info"]["initial_state"]
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    ENV_DIR.mkdir(parents=True, exist_ok=True)
    SNAPSHOT.write_text(json.dumps(task, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_tracker(state, DATA_DIR / "Safety_Incident_Tracker.xlsx")
    write_mbox(state["gmail"]["messages"], DATA_DIR / "Safety_Manual.mbox")
    (DATA_DIR / "Slack_Channels.json").write_text(json.dumps(state["slack"], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (DATA_DIR / "Workflow_Context.json").write_text(json.dumps({"meta":state["meta"]}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    for source in sorted(DATA_DIR.iterdir()):
        if source.is_file():
            shutil.copy2(source, ENV_DIR / source.name)


if __name__ == "__main__":
    main()
