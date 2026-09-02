#!/usr/bin/env python3
"""AutomationBench 5126 安全事故治理包的逐题参考解法。

这不是通用造题器。八起事故、三类内部路由、OSHA 触发、
工伤边界和公开 Slack 聚合均从本题四个冻结输入逐项重建。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import mailbox
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt as PptPt


INPUT_HASHES = {
    "Safety_Incident_Tracker.xlsx": "961eb03fb913283fc116d114793b15aa39277d5c043eccf5a48629e7341653f3",
    "Safety_Manual.mbox": "81e2472039256c090ac1fad2014d76c6590c5c4d10ec4bd0c70b82f0b009457b",
    "Slack_Channels.json": "6f4850069e074815ee1b160108497d37c4a1601cf2f37a15d65b11aef7fa9f92",
    "Workflow_Context.json": "76c4a0657bc27f2fd830b10613ec7ceae3adcdc5798aeee0f822d27bf69bbc99",
}
XLSX_NAME = "Safety_Incident_Audit.xlsx"
DOCX_NAME = "Safety_Incident_Control_Brief.docx"
PPTX_NAME = "Safety_Incident_Review.pptx"
MBOX_NAME = "Safety_Notification_Drafts.mbox"
UPDATES_NAME = "Safety_Tracker_Update_Drafts.json"
SLACK_NAME = "Safety_Slack_Summary_Draft.json"
EXECUTION_NAME = "Safety_Execution_Control.json"
FIXED_ISO_TIME = "2026-04-15T09:00:00Z"
FIXED_DATE = date(2026, 4, 15)
FIXED_ZIP_TIME = (2026, 4, 15, 9, 0, 0)

NAVY = "17365D"; BLUE = "5B9BD5"; GREEN = "E2F0D9"; AMBER = "FFF2CC"; RED = "F4CCCC"; GRAY = "E7E6E6"; WHITE = "FFFFFF"
THIN = Side(style="thin", color="B7C9D6")
PPT_NAVY = PptRGBColor(23, 54, 93); PPT_BLUE = PptRGBColor(91, 155, 213); PPT_LIGHT_BLUE = PptRGBColor(221, 235, 247)
PPT_GREEN = PptRGBColor(226, 240, 217); PPT_AMBER = PptRGBColor(255, 242, 204); PPT_RED = PptRGBColor(244, 204, 204); PPT_WHITE = PptRGBColor(255, 255, 255); PPT_GRAY = PptRGBColor(231, 230, 230)


@dataclass(frozen=True)
class IncidentRecord:
    source_row: int
    employee: str
    manager_email: str
    incident_date: date
    description: str
    severity: str
    treatment: str
    lost_days: str
    source_status: str
    notes: str


@dataclass(frozen=True)
class IncidentDecision:
    source: IncidentRecord
    hospitalization: str
    amputation: str
    vision_source: str
    osha_reportable: str
    proposed_status: str
    workers_comp_required: str
    legal_immediate: str
    recipients: str
    subject: str
    rationale: str


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def decode_message(message) -> str:
    if message.is_multipart():
        parts = []
        for part in message.walk():
            if part.get_content_type() == "text/plain":
                parts.append((part.get_payload(decode=True) or b"").decode(part.get_content_charset() or "utf-8", errors="replace"))
        return "\n".join(parts).strip()
    return (message.get_payload(decode=True) or b"").decode(message.get_content_charset() or "utf-8", errors="replace").strip()


def read_inputs(input_dir: Path) -> tuple[list[IncidentRecord], dict[str, str], dict, str]:
    for name, expected in INPUT_HASHES.items():
        source = input_dir / name
        if not source.is_file() or sha256(source) != expected:
            raise ValueError(f"输入 {name} 缺失或 SHA256 与冻结版本不一致")
    context = json.loads((input_dir / "Workflow_Context.json").read_text(encoding="utf-8"))
    if context != {"meta":{"schema_version":"0.1.0","current_time":FIXED_ISO_TIME}}:
        raise ValueError("冻结运行时间已变化")
    channels = json.loads((input_dir / "Slack_Channels.json").read_text(encoding="utf-8"))
    if channels != {"channels":[{"id":"C_SAF_5126","name":"safety","is_private":False}],"messages":[]}:
        raise ValueError("Safety Slack 快照已变化")

    workbook = load_workbook(input_dir / "Safety_Incident_Tracker.xlsx", read_only=True, data_only=True)
    if workbook.sheetnames != ["Recent Incidents"]:
        raise ValueError("安全事故台账工作表结构已变化")
    records: list[IncidentRecord] = []
    for row in workbook["Recent Incidents"].iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        raw_date = row[3]
        incident_date = raw_date.date() if isinstance(raw_date, datetime) else raw_date
        if not isinstance(incident_date, date):
            incident_date = date.fromisoformat(str(raw_date))
        records.append(IncidentRecord(int(row[0]),str(row[1]),str(row[2]),incident_date,str(row[4]),str(row[5]),str(row[6]),str(row[7]),str(row[8]),str(row[9] or "")))
    workbook.close()
    expected_records = [
        (2,"Ana Petrova","mgr-a@company.example.com","2026-04-14","Paper cut from filing cabinet","Minor","First aid - bandage applied","0","Pending",""),
        (3,"Ben Torres","mgr-b@company.example.com","2026-04-13","Slipped on wet floor, sprained ankle","Moderate","Urgent care visit, ankle brace","2","Pending",""),
        (4,"Carlos Mendoza","mgr-a@company.example.com","2026-04-12","Fell from loading dock, broken leg","Serious","Hospitalization - emergency surgery","TBD","Pending",""),
        (5,"Diana Osei","mgr-b@company.example.com","2026-04-14","Ergonomic discomfort - wrist pain","Minor","Ice pack, ergonomic assessment scheduled","0","Pending",""),
        (6,"Ethan Park","mgr-a@company.example.com","2026-04-11","Chemical splash in lab - eye injury","Serious","Hospitalization - potential loss of vision","5+","Pending",""),
        (7,"Fiona Strand","mgr-b@company.example.com","2026-04-14","Bumped head on low clearance doorway","Minor","Ice pack applied","0","Pending",""),
        (8,"George Tan","mgr-a@company.example.com","2026-04-10","Repetitive strain - back pain from lifting","Moderate","Physical therapy referral","1","Pending",""),
        (9,"Holly Ng","mgr-b@company.example.com","2026-04-13","Machine guard failure - finger amputation","Serious","Hospitalization - surgical amputation","TBD","Pending",""),
    ]
    actual = [(r.source_row,r.employee,r.manager_email,r.incident_date.isoformat(),r.description,r.severity,r.treatment,r.lost_days,r.source_status,r.notes) for r in records]
    if actual != expected_records:
        raise ValueError("八条冻结事故记录已变化")

    box = mailbox.mbox(input_dir / "Safety_Manual.mbox", create=False)
    messages = list(box); box.close()
    if len(messages) != 1:
        raise ValueError("Safety Manual MBOX 必须恰好一封")
    message = messages[0]
    manual = {
        "message_id":str(message.get("X-AutomationBench-Message-ID","")),
        "thread_id":str(message.get("X-AutomationBench-Thread-ID","")),
        "from":str(message.get("From","")), "to":str(message.get("To","")),
        "subject":str(message.get("Subject","")), "date":str(message.get("Date","")),
        "labels":str(message.get("X-AutomationBench-Labels","")), "body":decode_message(message),
    }
    if manual["message_id"] != "msg_safety_manual_5126" or manual["from"].casefold() != "safety-officer@company.example.com" or manual["subject"] != "Workplace Safety Incident Reporting Manual - 2026":
        raise ValueError("安全手册邮件身份已变化")
    required = ("Minor: First aid only","Moderate: Medical treatment needed","Serious: Hospitalization","Minor: Email the employee's manager","Moderate: Email the Safety Officer","Serious: Email the Safety Officer","workers' comp claim form","hospitalization, amputation, or loss of an eye","reported to OSHA within 24 hours","legal@company.example.com","retained for 5 years")
    if not all(token in manual["body"] for token in required):
        raise ValueError("安全手册关键证据不完整")
    return records, manual, channels, context["meta"]["current_time"]


def derive_decisions(records: list[IncidentRecord]) -> list[IncidentDecision]:
    decisions: list[IncidentDecision] = []
    for source in records:
        treatment = source.treatment.casefold(); combined = f"{source.description} {source.treatment}".casefold()
        hospitalization = "YES" if "hospitalization" in treatment else "NO"
        amputation = "YES" if "amputation" in combined else "NO"
        vision_source = "POTENTIAL_VISION_LOSS" if "potential loss of vision" in treatment else "NO"
        osha = "YES" if hospitalization == "YES" or amputation == "YES" else "NO"
        if source.severity == "Minor":
            status = "Documented"; workers = "NO"; legal = "NO"; recipients = source.manager_email
            subject = f"Safety Incident Documentation - {source.employee} - {source.incident_date.isoformat()}"
            rationale = "Minor: manager only; first-aid record; no workers' comp or OSHA trigger"
        elif source.severity == "Moderate":
            status = "Under Review"; workers = "YES"; legal = "NO"; recipients = f"safety-officer@company.example.com; {source.manager_email}"
            subject = f"Safety Incident Under Review - {source.employee} - {source.incident_date.isoformat()}"
            rationale = "Moderate: Safety Officer plus manager; workers' comp follow-up; no OSHA trigger"
        elif source.severity == "Serious":
            status = "Investigation"; workers = "YES"; legal = "YES" if osha == "YES" else "NO"
            base = f"safety-officer@company.example.com; {source.manager_email}; vp-ops@company.example.com"
            recipients = f"legal@company.example.com; {base}" if osha == "YES" else base
            subject = f"URGENT OSHA Review - {source.employee} - {source.incident_date.isoformat()}" if osha == "YES" else f"Serious Safety Incident - {source.employee} - {source.incident_date.isoformat()}"
            rationale = "Serious: standard internal routing plus immediate Legal priority because hospitalization is OSHA-reportable"
        else:
            raise ValueError(f"未知严重度：{source.severity}")
        decisions.append(IncidentDecision(source,hospitalization,amputation,vision_source,osha,status,workers,legal,recipients,subject,rationale))
    expected = {
        "Ana Petrova":("NO","NO","NO","NO","Documented","NO","NO","mgr-a@company.example.com"),
        "Ben Torres":("NO","NO","NO","NO","Under Review","YES","NO","safety-officer@company.example.com; mgr-b@company.example.com"),
        "Carlos Mendoza":("YES","NO","NO","YES","Investigation","YES","YES","legal@company.example.com; safety-officer@company.example.com; mgr-a@company.example.com; vp-ops@company.example.com"),
        "Diana Osei":("NO","NO","NO","NO","Documented","NO","NO","mgr-b@company.example.com"),
        "Ethan Park":("YES","NO","POTENTIAL_VISION_LOSS","YES","Investigation","YES","YES","legal@company.example.com; safety-officer@company.example.com; mgr-a@company.example.com; vp-ops@company.example.com"),
        "Fiona Strand":("NO","NO","NO","NO","Documented","NO","NO","mgr-b@company.example.com"),
        "George Tan":("NO","NO","NO","NO","Under Review","YES","NO","safety-officer@company.example.com; mgr-a@company.example.com"),
        "Holly Ng":("YES","YES","NO","YES","Investigation","YES","YES","legal@company.example.com; safety-officer@company.example.com; mgr-b@company.example.com; vp-ops@company.example.com"),
    }
    actual = {d.source.employee:(d.hospitalization,d.amputation,d.vision_source,d.osha_reportable,d.proposed_status,d.workers_comp_required,d.legal_immediate,d.recipients) for d in decisions}
    if actual != expected:
        raise ValueError("八起事故决策真值不一致")
    return decisions


def add_table(sheet, name: str) -> None:
    table = Table(displayName=name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)


def style_sheet(sheet, widths: dict[str, float], footer: str, *, a3: bool = False) -> None:
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions; sheet.sheet_view.showGridLines = False
    for column, width in widths.items(): sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY); cell.font = Font(name="Aptos", bold=True, color=WHITE, size=8); cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = Border(bottom=THIN)
    sheet.row_dimensions[1].height = 31
    grid = Side(style="hair", color="D9E2F3")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=7.6); cell.alignment = Alignment(vertical="top", wrap_text=True); cell.border = Border(left=grid,right=grid,bottom=grid)
    sheet.page_setup.orientation = "landscape"; sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if a3 else sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 1; sheet.sheet_properties.pageSetUpPr.fitToPage = True; sheet.print_options.horizontalCentered = True
    sheet.oddFooter.center.text = footer; sheet.print_area = sheet.dimensions


def normalize_ooxml(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="safety-5126-ooxml-") as temporary:
        target = Path(temporary) / path.name; axis_ids: dict[str,str] = {}
        with zipfile.ZipFile(path,"r") as source, zipfile.ZipFile(target,"w",zipfile.ZIP_DEFLATED,compresslevel=9) as destination:
            for name in sorted(source.namelist()):
                data = source.read(name)
                if name == "docProps/core.xml":
                    try:
                        root = etree.fromstring(data); ns={"dcterms":"http://purl.org/dc/terms/","cp":"http://schemas.openxmlformats.org/package/2006/metadata/core-properties"}
                        for tag in ("created","modified"):
                            for node in root.xpath(f"//dcterms:{tag}",namespaces=ns): node.text=FIXED_ISO_TIME
                        for tag in ("lastModifiedBy","revision"):
                            for node in root.xpath(f"//cp:{tag}",namespaces=ns): node.text="EvoCowork" if tag=="lastModifiedBy" else "1"
                        data=etree.tostring(root,xml_declaration=True,encoding="UTF-8",standalone=True)
                    except etree.XMLSyntaxError: pass
                if name.startswith("xl/charts/chart") and name.endswith(".xml"):
                    try:
                        root=etree.fromstring(data); ns={"c":"http://schemas.openxmlformats.org/drawingml/2006/chart"}
                        for node in root.xpath("//c:axId",namespaces=ns):
                            old=node.get("val")
                            if old and old not in axis_ids: axis_ids[old]=str(51_260_001+len(axis_ids))
                        for node in root.xpath("//c:axId | //c:crossAx",namespaces=ns):
                            old=node.get("val")
                            if old in axis_ids: node.set("val",axis_ids[old])
                        data=etree.tostring(root,xml_declaration=True,encoding="UTF-8",standalone=True)
                    except etree.XMLSyntaxError: pass
                info=zipfile.ZipInfo(name,FIXED_ZIP_TIME); info.compress_type=zipfile.ZIP_DEFLATED; info.external_attr=0o600<<16; destination.writestr(info,data)
        shutil.copy2(target,path)


def recalculate_xlsx(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="safety-5126-recalc-") as temporary:
        root=Path(temporary); source_dir=root/"source"; output_dir=root/"output"; profile_dir=root/"profile"; source_dir.mkdir(); output_dir.mkdir(); source=source_dir/path.name; shutil.copy2(path,source)
        process=subprocess.run(["libreoffice","--headless",f"-env:UserInstallation=file://{profile_dir}","--convert-to","xlsx","--outdir",str(output_dir),str(source)],stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,timeout=300)
        converted=output_dir/path.name
        if process.returncode!=0 or not converted.is_file(): raise RuntimeError(f"LibreOffice 公式重算失败：{process.stdout[-1000:]}")
        shutil.copy2(converted,path)


def source_trace(row: int) -> str:
    return f"Safety_Incident_Tracker.xlsx:Recent Incidents row {row} | Safety_Manual.mbox:msg_safety_manual_5126 | Workflow_Context.json:meta.current_time"


POLICY_ROWS = [
    ("P001","Classification","Minor","First aid only; no lost work time","Safety_Manual.mbox:Chapter 2"),
    ("P002","Classification","Moderate","Medical treatment; fewer than 3 lost work days","Safety_Manual.mbox:Chapter 2"),
    ("P003","Classification","Serious","Hospitalization, more than 3 lost work days, or restricted duty","Safety_Manual.mbox:Chapter 2"),
    ("P004","Internal routing","Minor","Manager only; Status=Documented","Safety_Manual.mbox:Chapter 3"),
    ("P005","Internal routing","Moderate","Safety Officer + manager; Status=Under Review","Safety_Manual.mbox:Chapter 3"),
    ("P006","Internal routing","Serious","Safety Officer + manager + VP Operations; Status=Investigation","Safety_Manual.mbox:Chapter 3"),
    ("P007","Workers' compensation","Moderate/Serious","Claim form required; Benefits handles separately","Safety_Manual.mbox:Chapter 4"),
    ("P008","OSHA","Hospitalization/amputation/loss of eye","Legal immediately; OSHA within 24 hours; precedence","Safety_Manual.mbox:Chapter 5"),
    ("P009","Return to work","Return from injury leave","Assessment required","Safety_Manual.mbox:Chapter 6"),
    ("P010","Retention","All incidents","Retain records for 5 years","Safety_Manual.mbox:Chapter 7"),
]


def write_workbook(path: Path, records: list[IncidentRecord], manual: dict[str,str], channels: dict, frozen_time: str) -> None:
    workbook=Workbook(); workbook.remove(workbook.active); workbook.calculation.fullCalcOnLoad=True; workbook.calculation.forceFullCalc=True; workbook.calculation.calcMode="auto"
    workbook.properties.creator="EvoCowork / AutomationBench 5126"; workbook.properties.title="Safety Incident Audit"; workbook.properties.subject="Source-backed offline safety incident routing and OSHA control"; workbook.properties.description="Eight frozen incidents; no external action executed."

    source=workbook.create_sheet("Source_Incidents"); source.append(["Source Row","Employee","Manager Email","Incident Date","Description","Severity","Treatment","Lost Days","Status","Notes"])
    for r in records: source.append([r.source_row,r.employee,r.manager_email,r.incident_date,r.description,r.severity,r.treatment,r.lost_days,r.source_status,r.notes])
    for row in range(2,10): source.cell(row,4).number_format="yyyy-mm-dd"; source.row_dimensions[row].height=38
    add_table(source,"SafetySourceIncidents"); style_sheet(source,{"A":11,"B":20,"C":29,"D":16,"E":39,"F":13,"G":41,"H":12,"I":13,"J":15},"Frozen AutomationBench 5126 source incidents",a3=True)

    source_manual=workbook.create_sheet("Source_Manual"); source_manual.append(["Message ID","Thread ID","From","To","Subject","Date","Labels","Body","Frozen Time","Source"]); source_manual.append([manual["message_id"],manual["thread_id"],manual["from"],manual["to"],manual["subject"],manual["date"],manual["labels"],manual["body"],frozen_time,"Safety_Manual.mbox"]); source_manual.row_dimensions[2].height=250
    add_table(source_manual,"SafetySourceManual"); style_sheet(source_manual,{"A":25,"B":24,"C":29,"D":29,"E":42,"F":27,"G":12,"H":110,"I":28,"J":24},"Frozen safety manual and runtime",a3=True)

    policy=workbook.create_sheet("Policy_Source"); policy.append(["Policy ID","Chapter","Applies To","Requirement","Source Trace"])
    for row in POLICY_ROWS: policy.append(row)
    add_table(policy,"SafetyPolicySource"); style_sheet(policy,{"A":12,"B":24,"C":35,"D":75,"E":38},"Ten clauses reconstructed from the frozen manual",a3=True)

    evidence=workbook.create_sheet("Evidence_Register"); evidence.append(["Evidence ID","Type","Entity/Control","Value","Source Trace"])
    for index,r in enumerate(records,1): evidence.append((f"E{index:03d}","SOURCE_INCIDENT",r.employee,f"{r.incident_date.isoformat()} | {r.severity} | {r.treatment}",f"Safety_Incident_Tracker.xlsx:Recent Incidents row {r.source_row}"))
    evidence_rows=[
        ("E009","MANUAL","Reporting manual",manual["subject"],"Safety_Manual.mbox:msg_safety_manual_5126"),
        ("E010","TIME","Frozen time",frozen_time,"Workflow_Context.json:meta.current_time"),
        ("E011","CHANNEL","Safety channel","C_SAF_5126 | safety | public","Slack_Channels.json:C_SAF_5126"),
        ("E012","CONTROL","Minor count","3","Incident_Decisions:Severity"),
        ("E013","CONTROL","Moderate count","2","Incident_Decisions:Severity"),
        ("E014","CONTROL","Serious count","3","Incident_Decisions:Severity"),
        ("E015","CONTROL","OSHA/Legal immediate count","3","Incident_Decisions:OSHA Reportable"),
        ("E016","CONTROL","Workers comp follow-up count","5","Incident_Decisions:Workers Comp Required"),
        ("E017","CONTROL","Email draft count","8","Communication_Register"),
        ("E018","CONTROL","Tracker update draft count","8","Tracker update payload"),
        ("E019","CONTROL","Slack aggregate draft count","1","Slack summary payload"),
        ("E020","CONTROL","Emails sent","0","Execution control"),
        ("E021","CONTROL","Tracker updates applied","0","Execution control"),
        ("E022","CONTROL","Real side effects","0","Execution control"),
    ]
    for row in evidence_rows: evidence.append(row)
    add_table(evidence,"SafetyEvidenceRegister"); style_sheet(evidence,{"A":13,"B":20,"C":34,"D":70,"E":50},"Twenty-two source and control evidence rows",a3=True)

    decisions=workbook.create_sheet("Incident_Decisions")
    headers=["Source Row","Employee","Manager Email","Incident Date","Description","Severity","Treatment","Lost Days","Source Status","Frozen Time","Is Minor","Is Moderate","Is Serious","Hospitalization Trigger","Amputation Trigger","Vision Source","OSHA Reportable","Proposed Status","Workers Comp Required","Legal Immediate","Recipients","Email Subject","Slack Scope","Tracker Update Draft","Email State","Tracker Applied","Slack Posted","OSHA Filed","Retention Years","Policy Basis","Source Trace"]
    decisions.append(headers)
    for row in range(2,10):
        src=row
        formulas=[
            f"=INDEX(Source_Incidents!$A$2:$A$9,ROW()-1)",f"=INDEX(Source_Incidents!$B$2:$B$9,ROW()-1)",f"=INDEX(Source_Incidents!$C$2:$C$9,ROW()-1)",f"=INDEX(Source_Incidents!$D$2:$D$9,ROW()-1)",f"=INDEX(Source_Incidents!$E$2:$E$9,ROW()-1)",f"=INDEX(Source_Incidents!$F$2:$F$9,ROW()-1)",f"=INDEX(Source_Incidents!$G$2:$G$9,ROW()-1)",f"=INDEX(Source_Incidents!$H$2:$H$9,ROW()-1)",f"=INDEX(Source_Incidents!$I$2:$I$9,ROW()-1)","=Evidence_Register!D11",
            f'=IF(F{row}="Minor","YES","NO")',f'=IF(F{row}="Moderate","YES","NO")',f'=IF(F{row}="Serious","YES","NO")',f'=IF(ISNUMBER(SEARCH("hospitalization",G{row})),"YES","NO")',f'=IF(ISNUMBER(SEARCH("amputation",E{row}&" "&G{row})),"YES","NO")',f'=IF(ISNUMBER(SEARCH("potential loss of vision",G{row})),"POTENTIAL_VISION_LOSS","NO")',f'=IF(OR(N{row}="YES",O{row}="YES"),"YES","NO")',f'=IF(K{row}="YES","Documented",IF(L{row}="YES","Under Review",IF(M{row}="YES","Investigation","INVALID")))',f'=IF(OR(L{row}="YES",M{row}="YES"),"YES","NO")',f'=IF(Q{row}="YES","YES","NO")',f'=IF(K{row}="YES",C{row},IF(L{row}="YES","safety-officer@company.example.com; "&C{row},IF(Q{row}="YES","legal@company.example.com; safety-officer@company.example.com; "&C{row}&"; vp-ops@company.example.com","safety-officer@company.example.com; "&C{row}&"; vp-ops@company.example.com")))',f'=IF(K{row}="YES","Safety Incident Documentation - "&B{row}&" - "&TEXT(D{row},"yyyy-mm-dd"),IF(L{row}="YES","Safety Incident Under Review - "&B{row}&" - "&TEXT(D{row},"yyyy-mm-dd"),IF(Q{row}="YES","URGENT OSHA Review - "&B{row}&" - "&TEXT(D{row},"yyyy-mm-dd"),"Serious Safety Incident - "&B{row}&" - "&TEXT(D{row},"yyyy-mm-dd"))))',
            '="AGGREGATE_ONLY"','="YES"','="DRAFT_NOT_SENT"','="NO"','="NO"','="NO"','=5',f'=IF(K{row}="YES","P004,P010",IF(L{row}="YES","P005,P007,P010",IF(Q{row}="YES","P006,P007,P008,P010","P006,P007,P010")))',f'="Safety_Incident_Tracker.xlsx:Recent Incidents row "&A{row}&" | Safety_Manual.mbox:msg_safety_manual_5126 | Workflow_Context.json:meta.current_time"'
        ]
        decisions.append(formulas); decisions.cell(row,4).number_format="yyyy-mm-dd"; decisions.row_dimensions[row].height=50
    fills={"$F2=\"Minor\"":GREEN,"$F2=\"Moderate\"":AMBER,"$F2=\"Serious\"":RED,"$Q2=\"YES\"":"F8CBAD"}
    for formula_text,color in fills.items(): decisions.conditional_formatting.add("A2:AE9",FormulaRule(formula=[formula_text],fill=PatternFill("solid",fgColor=color)))
    add_table(decisions,"SafetyIncidentDecisions"); style_sheet(decisions,{chr(64+i):12 for i in range(1,27)}|{"AA":12,"AB":12,"AC":12,"AD":36,"AE":54},"Formula-derived incident routing",a3=True)

    def create_route_sheet(name: str, table_name: str, decision_rows: list[int], columns: list[tuple[str,int]], footer: str) -> None:
        sheet=workbook.create_sheet(name); sheet.append([label for label,_ in columns])
        for source_row in decision_rows:
            sheet.append([f"=Incident_Decisions!{chr(64+column) if column<=26 else 'A'+chr(64+column-26)}{source_row}" for _,column in columns])
        add_table(sheet,table_name); style_sheet(sheet,{chr(64+i):22 for i in range(1,len(columns)+1)},footer,a3=True)

    create_route_sheet("Minor_Routing","SafetyMinorRouting",[2,5,7],[
        ("Source Row",1),("Employee",2),("Incident Date",4),("Severity",6),("Proposed Status",18),("Recipients",21),("Email Subject",22),("OSHA Reportable",17),("Workers Comp",19),("Email State",25),("Tracker Applied",26),("Retention Years",29),("Policy Basis",30),("Source Trace",31)],"Minor: manager only")
    create_route_sheet("Moderate_Routing","SafetyModerateRouting",[3,8],[
        ("Source Row",1),("Employee",2),("Incident Date",4),("Description",5),("Treatment",7),("Lost Days",8),("Severity",6),("Proposed Status",18),("Recipients",21),("Workers Comp",19),("OSHA Reportable",17),("Legal Immediate",20),("Email State",25),("Tracker Applied",26),("Policy Basis",30),("Source Trace",31)],"Moderate: Safety Officer plus manager")
    create_route_sheet("Serious_OSHA_Routing","SafetySeriousRouting",[4,6,9],[
        ("Source Row",1),("Employee",2),("Manager Email",3),("Incident Date",4),("Description",5),("Treatment",7),("Lost Days",8),("Hospitalization",14),("Amputation",15),("Vision Source",16),("OSHA Reportable",17),("Legal Immediate",20),("Proposed Status",18),("Workers Comp",19),("Recipients",21),("Email Subject",22),("Email State",25),("OSHA Filed",28),("Policy Basis",30),("Source Trace",31)],"Serious and OSHA/Legal priority")

    communication=workbook.create_sheet("Communication_Register"); communication.append(["Order","Source Row","Employee","Severity","Communication Type","Recipients","Subject","OSHA Reportable","Workers Comp","Tracker Status","Email State","Sent","Attachment Count","Source Trace"])
    for order,row in enumerate(range(2,10),1):
        communication.append([f"={order}",f"=Incident_Decisions!A{row}",f"=Incident_Decisions!B{row}",f"=Incident_Decisions!F{row}",f'=IF(Incident_Decisions!Q{row}="YES","OSHA_LEGAL_AND_INTERNAL",IF(Incident_Decisions!K{row}="YES","MANAGER_ONLY","SAFETY_AND_INTERNAL"))',f"=Incident_Decisions!U{row}",f"=Incident_Decisions!V{row}",f"=Incident_Decisions!Q{row}",f"=Incident_Decisions!S{row}",f"=Incident_Decisions!R{row}",'="DRAFT_NOT_SENT"','="NO"','=0',f"=Incident_Decisions!AE{row}"])
    add_table(communication,"SafetyCommunicationRegister"); style_sheet(communication,{"A":9,"B":11,"C":20,"D":13,"E":30,"F":75,"G":52,"H":17,"I":16,"J":18,"K":20,"L":10,"M":14,"N":60},"Eight ordered offline email drafts",a3=True)

    slack=workbook.create_sheet("Slack_Summary_Control"); slack.append(["Channel ID","Channel","Public","Draft Count","Posted Count","Incident Count","Minor","Moderate","Serious","OSHA/Legal Review","Workers Comp","Contains Names","Contains Medical Details","State","Draft Text","Source Trace"])
    slack.append(["C_SAF_5126","safety","YES",1,0,8,3,2,3,3,5,"NO","NO","DRAFT_NOT_POSTED",slack_text(),"Slack_Channels.json:C_SAF_5126 | Incident_Decisions aggregate"])
    add_table(slack,"SafetySlackControl"); style_sheet(slack,{"A":18,"B":14,"C":10,"D":12,"E":12,"F":14,"G":10,"H":12,"I":10,"J":20,"K":16,"L":16,"M":23,"N":21,"O":105,"P":58},"Public channel aggregate only",a3=True)

    summary=workbook.create_sheet("Control_Summary"); summary.append(["Metric","Actual","Expected","Status","Evidence"])
    metrics=[
        ("Source incidents",'=COUNTA(Source_Incidents!B2:B9)',8,"All frozen rows"),("Minor",'=COUNTIF(Incident_Decisions!F2:F9,"Minor")',3,"Ana; Diana; Fiona"),("Moderate",'=COUNTIF(Incident_Decisions!F2:F9,"Moderate")',2,"Ben; George"),("Serious",'=COUNTIF(Incident_Decisions!F2:F9,"Serious")',3,"Carlos; Ethan; Holly"),("Documented",'=COUNTIF(Incident_Decisions!R2:R9,"Documented")',3,"Minor status"),("Under Review",'=COUNTIF(Incident_Decisions!R2:R9,"Under Review")',2,"Moderate status"),("Investigation",'=COUNTIF(Incident_Decisions!R2:R9,"Investigation")',3,"Serious status"),("OSHA/Legal immediate",'=COUNTIF(Incident_Decisions!Q2:Q9,"YES")',3,"Three hospitalizations"),("Workers comp follow-up",'=COUNTIF(Incident_Decisions!S2:S9,"YES")',5,"Moderate + Serious"),("Email drafts",'=COUNTA(Communication_Register!C2:C9)',8,"One per incident"),("Emails sent",'=COUNTIF(Communication_Register!L2:L9,"YES")',0,"Offline only"),("Tracker update drafts",'=COUNTIF(Incident_Decisions!X2:X9,"YES")',8,"One per incident"),("Tracker updates applied",'=COUNTIF(Incident_Decisions!Z2:Z9,"YES")',0,"Source unchanged"),("Slack drafts",'=Slack_Summary_Control!D2',1,"Aggregate only"),("Slack posts",'=Slack_Summary_Control!E2',0,"Offline only"),("OSHA reports filed",'=COUNTIF(Incident_Decisions!AB2:AB9,"YES")',0,"Legal draft only"),("Workers comp forms created",'=0',0,"Benefits handles separately"),("Exact OSHA deadlines asserted",'=0',0,"Incident time unavailable"),("Names in Slack",'=COUNTIF(Slack_Summary_Control!L2,"YES")',0,"Public aggregate"),("Medical details in Slack",'=COUNTIF(Slack_Summary_Control!M2,"YES")',0,"Public aggregate"),("Retention years",'=MIN(Incident_Decisions!AC2:AC9)',5,"All records"),("Real side effects",'=0',0,"Controlled draft only")]
    for index,(metric,formula_value,expected_value,evidence_text) in enumerate(metrics,2): summary.append([metric,formula_value,expected_value,f'=IF(B{index}=C{index},"PASS","FAIL")',evidence_text])
    chart=BarChart(); chart.title="Safety routing distribution"; chart.height=6.5; chart.width=12; chart.add_data(Reference(summary,min_col=2,min_row=2,max_row=4),titles_from_data=False); chart.set_categories(Reference(summary,min_col=1,min_row=2,max_row=4)); chart.legend=None; summary.add_chart(chart,"G2")
    add_table(summary,"SafetyControlSummary"); style_sheet(summary,{"A":31,"B":14,"C":14,"D":13,"E":45,"G":15},"Twenty-two count and control checks",a3=True)

    workbook.save(path); recalculate_xlsx(path); normalize_ooxml(path)


def set_cell_shading(cell, color: str) -> None:
    properties=cell._tc.get_or_add_tcPr(); shading=properties.find(qn("w:shd"))
    if shading is None: shading=OxmlElement("w:shd"); properties.append(shading)
    shading.set(qn("w:fill"),color)


def add_doc_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float]|None=None) -> None:
    table=document.add_table(rows=1,cols=len(headers)); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style="Table Grid"
    for i,header in enumerate(headers):
        cell=table.rows[0].cells[i]; cell.text=header; set_cell_shading(cell,NAVY)
        for run in cell.paragraphs[0].runs: run.font.name="Carlito"; run.font.color.rgb=RGBColor(255,255,255); run.font.bold=True; run.font.size=Pt(7)
    for row in rows:
        cells=table.add_row().cells
        for i,value in enumerate(row):
            cells[i].text=str(value); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cells[i].paragraphs:
                paragraph.paragraph_format.space_after=Pt(0)
                for run in paragraph.runs: run.font.name="Carlito"; run.font.size=Pt(6.5)
        if widths:
            for i,width in enumerate(widths): cells[i].width=Cm(width)
    document.add_paragraph().paragraph_format.space_after=Pt(1)


def configure_document(document: Document, title: str) -> None:
    section=document.sections[0]; section.orientation=WD_ORIENT.LANDSCAPE; section.page_width=Cm(29.7); section.page_height=Cm(21); section.top_margin=Cm(1); section.bottom_margin=Cm(1); section.left_margin=Cm(1); section.right_margin=Cm(1)
    styles=document.styles; styles["Normal"].font.name="Carlito"; styles["Normal"].font.size=Pt(8); styles["Title"].font.name="Carlito"; styles["Title"].font.size=Pt(21); styles["Title"].font.color.rgb=RGBColor(23,54,93); styles["Heading 1"].font.name="Carlito"
    document.core_properties.title=title; document.core_properties.author="EvoCowork"; document.core_properties.subject="AutomationBench 5126 safety incident control package"


def write_brief(path: Path, decisions: list[IncidentDecision]) -> None:
    document=Document(); configure_document(document,"Safety Incident Control Brief")
    title=document.add_paragraph("Safety Incident Control Brief",style="Title"); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=document.add_paragraph("Frozen at 2026-04-15 09:00 UTC · 8 incidents · 3 Minor · 2 Moderate · 3 Serious · 3 OSHA/Legal immediate · 5 workers' comp follow-ups · all actions remain offline drafts."); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    add_doc_table(document,["Severity","Count","Internal route","Tracker status","Workers' comp","OSHA/Legal"],[
        ["Minor","3","Manager only","Documented","No","No trigger"],
        ["Moderate","2","Safety Officer + manager","Under Review","Required / Benefits","No trigger"],
        ["Serious","3","Legal first + Safety + manager + VP Ops","Investigation","Required / Benefits","3 hospitalization triggers"],
    ])
    add_doc_table(document,["Row","Employee","Date","Severity","Treatment / trigger","Status","Recipients","OSHA","WC"],[[str(d.source.source_row),d.source.employee,d.source.incident_date.isoformat(),d.source.severity,d.source.treatment,d.proposed_status,d.recipients,d.osha_reportable,d.workers_comp_required] for d in decisions])
    document.add_page_break()
    document.add_heading("OSHA, Legal, and evidence boundaries",level=1)
    document.add_paragraph("Carlos Mendoza, Ethan Park, and Holly Ng are OSHA-reportable because each frozen treatment explicitly records hospitalization. Holly also has a finger amputation; Ethan's source records potential loss of vision, not a confirmed loss of an eye. Legal must be notified immediately while Serious internal routing remains in force.")
    add_doc_table(document,["Employee","Source trigger","Legal","Standard Serious route","Deadline statement","Execution"],[
        ["Carlos Mendoza","Hospitalization - emergency surgery","Immediate draft","Safety + mgr-a + VP Ops","Exact incident time unavailable; no exact 24h cutoff asserted","OSHA filed 0"],
        ["Ethan Park","Hospitalization; potential loss of vision","Immediate draft","Safety + mgr-a + VP Ops","Exact incident time unavailable; no exact 24h cutoff asserted","OSHA filed 0"],
        ["Holly Ng","Hospitalization; finger amputation","Immediate draft","Safety + mgr-b + VP Ops","Exact incident time unavailable; no exact 24h cutoff asserted","OSHA filed 0"],
    ])
    document.add_paragraph("Workers' compensation follows Chapter 4: Ben, Carlos, Ethan, George, and Holly require follow-up, but the Benefits team handles claim forms separately. This package creates 0 workers' comp forms and does not claim submission. All eight records require five-year retention. Return-to-work assessment applies only if an employee later returns from injury leave; no such return event appears in the frozen inputs.")
    add_doc_table(document,["Evidence","Frozen fact","Boundary"],[
        ["Safety manual","Chapters 2-7 retained verbatim","No new policy invented"],
        ["Incident tracker","Eight source rows preserved","No medical diagnosis inferred"],
        ["Frozen time","2026-04-15T09:00:00Z","Incident dates lack times"],
        ["Public #safety","Aggregate draft only","No names, manager emails, treatments, or descriptions"],
    ])
    document.add_page_break()
    document.add_heading("Communication and execution control",level=1)
    add_doc_table(document,["Order","Employee","Draft subject","Recipients","State"],[[str(i),d.source.employee,d.subject,d.recipients,"DRAFT_NOT_SENT"] for i,d in enumerate(decisions,1)])
    add_doc_table(document,["Control","Draft/proposed","Executed","State"],[
        ["Incident email","8","0","DRAFT_NOT_SENT"],["Tracker update","8","0","SOURCE_UNCHANGED"],["Slack aggregate summary","1","0","DRAFT_NOT_POSTED"],["OSHA filing","3 Legal reviews","0","NOT_FILED"],["Workers' comp forms","5 follow-ups","0","BENEFITS_SEPARATE"],["Real side effects","0","0","OFFLINE_ONLY"],
    ])
    document.add_paragraph("Public #safety draft: 8 reviewed; 3 minor, 2 moderate, 3 serious; 3 require immediate Legal/OSHA review; 5 require workers' compensation follow-up. Employee names, manager addresses, medical details, and incident descriptions are excluded.")
    document.save(path); normalize_ooxml(path)


def ppt_textbox(slide,left,top,width,height,text,size=12,bold=False,color=PPT_NAVY,fill=None,align=PP_ALIGN.LEFT):
    shape=slide.shapes.add_textbox(Inches(left),Inches(top),Inches(width),Inches(height)) if fill is None else slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(left),Inches(top),Inches(width),Inches(height))
    if fill is not None: shape.fill.solid(); shape.fill.fore_color.rgb=fill; shape.line.color.rgb=fill
    frame=shape.text_frame; frame.clear(); frame.margin_left=Inches(.08); frame.margin_right=Inches(.08); frame.margin_top=Inches(.04); frame.margin_bottom=Inches(.04); frame.vertical_anchor=MSO_ANCHOR.MIDDLE
    paragraph=frame.paragraphs[0]; paragraph.text=text; paragraph.alignment=align
    for run in paragraph.runs: run.font.name="Aptos"; run.font.size=PptPt(size); run.font.bold=bold; run.font.color.rgb=color
    return shape


def ppt_table(slide,left,top,width,height,headers,rows,col_widths=None,font_size=8):
    shape=slide.shapes.add_table(len(rows)+1,len(headers),Inches(left),Inches(top),Inches(width),Inches(height)); table=shape.table
    if col_widths:
        for i,value in enumerate(col_widths): table.columns[i].width=Inches(value)
    for row_index,row in enumerate([headers]+rows):
        for column_index,value in enumerate(row):
            cell=table.cell(row_index,column_index); cell.text=str(value); cell.margin_left=Inches(.035); cell.margin_right=Inches(.035); cell.margin_top=Inches(.02); cell.margin_bottom=Inches(.02); cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            cell.fill.solid(); cell.fill.fore_color.rgb=PPT_NAVY if row_index==0 else (PPT_LIGHT_BLUE if row_index%2 else PPT_WHITE)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment=PP_ALIGN.CENTER
                for run in paragraph.runs: run.font.name="Aptos"; run.font.size=PptPt(font_size); run.font.bold=row_index==0; run.font.color.rgb=PPT_WHITE if row_index==0 else PPT_NAVY
    return shape


def write_presentation(path: Path, decisions: list[IncidentDecision]) -> None:
    presentation=Presentation(); presentation.slide_width=Inches(13.333); presentation.slide_height=Inches(7.5); presentation.core_properties.title="Safety Incident Review"; presentation.core_properties.author="EvoCowork"; blank=presentation.slide_layouts[6]
    slide=presentation.slides.add_slide(blank); ppt_textbox(slide,.45,.25,12.4,.55,"Safety incident routing · frozen 2026-04-15",24,True)
    ppt_textbox(slide,.45,.9,12.4,.48,"8 reviewed · 3 Minor · 2 Moderate · 3 Serious · 3 OSHA/Legal immediate · 5 workers' comp follow-ups",13,True,PPT_WHITE,PPT_BLUE,PP_ALIGN.CENTER)
    ppt_table(slide,.55,1.55,12.2,3.05,["Severity","Count","Internal route","Status","Workers' comp","OSHA / Legal"],[
        ["Minor","3","Manager only","Documented","No","No"],["Moderate","2","Safety Officer + manager","Under Review","Required / Benefits","No"],["Serious","3","Legal first + Safety + manager + VP Ops","Investigation","Required / Benefits","3 hospitalization triggers"],
    ],[1.2,0.8,4.0,1.55,2.0,2.65],10)
    ppt_textbox(slide,.55,4.9,6.0,.72,"Exact OSHA cutoff time cannot be computed: source has dates, not incident times.",12,True,PPT_NAVY,PPT_AMBER)
    ppt_textbox(slide,6.75,4.9,6.0,.72,"Offline control: emails sent 0 · updates applied 0 · Slack posts 0 · OSHA filed 0.",12,True,PPT_NAVY,PPT_GREEN)
    ppt_textbox(slide,.55,5.85,12.2,.72,"Public #safety: aggregate draft only; no employee names, manager addresses, treatments, or incident descriptions.",12,True,PPT_WHITE,PPT_NAVY,PP_ALIGN.CENTER)

    slide=presentation.slides.add_slide(blank); ppt_textbox(slide,.45,.25,12.4,.55,"Eight source-backed decisions",24,True)
    rows=[]
    for d in decisions: rows.append([str(d.source.source_row),d.source.employee,d.source.severity,d.source.treatment,d.proposed_status,d.osha_reportable,d.workers_comp_required,d.recipients])
    ppt_table(slide,.35,.95,12.65,5.65,["Row","Employee","Severity","Treatment / trigger","Status","OSHA","WC","Recipients"],rows,[.55,1.4,.8,2.65,1.15,.55,.45,5.1],6.5)
    ppt_textbox(slide,.45,6.77,12.4,.36,"Carlos: hospitalization · Ethan: hospitalization + potential vision loss · Holly: hospitalization + finger amputation",9,True,PPT_WHITE,PPT_NAVY,PP_ALIGN.CENTER)

    slide=presentation.slides.add_slide(blank); ppt_textbox(slide,.45,.25,12.4,.55,"Draft queue and zero-side-effect controls",24,True)
    ppt_table(slide,.4,.95,12.55,4.55,["#","Employee","Type","Recipients","State"],[[str(i),d.source.employee,"OSHA_LEGAL_AND_INTERNAL" if d.osha_reportable=="YES" else ("MANAGER_ONLY" if d.source.severity=="Minor" else "SAFETY_AND_INTERNAL"),d.recipients,"DRAFT_NOT_SENT"] for i,d in enumerate(decisions,1)],[.45,1.4,2.25,6.65,1.8],7)
    ppt_textbox(slide,.45,5.7,3.95,.72,"Email 8 drafts · 0 sent\nTracker 8 drafts · 0 applied",11,True,PPT_NAVY,PPT_LIGHT_BLUE,PP_ALIGN.CENTER)
    ppt_textbox(slide,4.68,5.7,3.95,.72,"Slack 1 draft · 0 posted\nNames / medical details = 0",11,True,PPT_NAVY,PPT_GREEN,PP_ALIGN.CENTER)
    ppt_textbox(slide,8.91,5.7,3.95,.72,"OSHA reports filed = 0\nWorkers' comp forms = 0",11,True,PPT_NAVY,PPT_AMBER,PP_ALIGN.CENTER)
    ppt_textbox(slide,.45,6.62,12.4,.4,"Source tracker unchanged · no external messages · no filing · no real side effects",10,True,PPT_WHITE,PPT_NAVY,PP_ALIGN.CENTER)
    presentation.save(path); normalize_ooxml(path)


def addresses(decision: IncidentDecision) -> list[str]:
    return [item.strip() for item in decision.recipients.split(";") if item.strip()]


def email_body(decision: IncidentDecision) -> str:
    source=decision.source
    if decision.osha_reportable=="YES": osha=f"YES — {source.treatment}. Notify Legal immediately; exact incident time is unavailable, so no exact 24-hour cutoff is asserted. OSHA filing completed: NO."
    else: osha="NO source trigger. Legal is not included."
    workers="REQUIRED — Benefits team handles the claim form separately; form created: NO." if decision.workers_comp_required=="YES" else "NOT REQUIRED by Chapter 4 for Minor incidents."
    return "\n".join([
        "SAFETY INCIDENT CONTROL DRAFT","",f"Employee: {source.employee}",f"Source Row: {source.source_row}",f"Incident Date: {source.incident_date.isoformat()}",f"Severity: {source.severity}",f"Description: {source.description}",f"Treatment: {source.treatment}",f"Lost Days: {source.lost_days}",f"Proposed Tracker Status: {decision.proposed_status}",f"Internal Routing: {decision.recipients}",f"OSHA / Legal: {osha}",f"Workers' Compensation: {workers}","Retention: 5 years.","","DRAFT_NOT_SENT — email sent: 0; tracker updated: 0; Slack posted: 0; real side effects: 0.","- HR Ops",
    ])


def write_mbox(path: Path, decisions: list[IncidentDecision]) -> None:
    chunks: list[bytes]=[]
    for order,decision in enumerate(decisions,1):
        message=EmailMessage(policy=SMTP); message["From"]="hr-ops@company.example.com"; message["To"] = ", ".join(addresses(decision)); message["Subject"]=decision.subject; message["Date"]="Wed, 15 Apr 2026 09:00:00 +0000"; message["Message-ID"]=f"<safety-incident-5126-{order}@evocowork.local>"; message["X-EvoCowork-Status"]="DRAFT_NOT_SENT"; message["X-EvoCowork-Order"]=str(order); message["X-EvoCowork-Source-Row"]=str(decision.source.source_row); message["X-EvoCowork-Severity"]=decision.source.severity; message["X-EvoCowork-Proposed-Status"]=decision.proposed_status; message["X-EvoCowork-OSHA-Reportable"]=decision.osha_reportable; message["X-EvoCowork-Workers-Comp-Required"]=decision.workers_comp_required; message["X-EvoCowork-Email-Sent"]="NO"; message["X-EvoCowork-Tracker-Applied"]="NO"; message["X-EvoCowork-Slack-Posted"]="NO"; message["X-EvoCowork-OSHA-Filed"]="NO"; message.set_content(email_body(decision),charset="utf-8")
        buffer=BytesIO(); BytesGenerator(buffer,policy=SMTP,mangle_from_=True,maxheaderlen=0).flatten(message); chunks.append(f"From hr-ops@company.example.com Wed Apr 15 09:00:00 2026\n".encode("ascii")); chunks.append(buffer.getvalue().replace(b"\r\n",b"\n")); chunks.append(b"\n")
    path.write_bytes(b"".join(chunks))


def updates_payload(decisions: list[IncidentDecision]) -> dict:
    entries=[]
    for decision in decisions:
        source=decision.source
        entries.append({"source_row":source.source_row,"employee":source.employee,"source_status":source.source_status,"severity":source.severity,"proposed_validation_status":decision.proposed_status,"hospitalization_trigger":decision.hospitalization=="YES","amputation_trigger":decision.amputation=="YES","vision_source":decision.vision_source,"osha_reportable":decision.osha_reportable=="YES","legal_immediate":decision.legal_immediate=="YES","workers_comp_required":decision.workers_comp_required=="YES","recipients":addresses(decision),"email_subject":decision.subject,"email_state":"DRAFT_NOT_SENT","tracker_update_applied":False,"osha_filed":False,"workers_comp_form_created":False,"slack_posted":False,"source_trace":[f"Safety_Incident_Tracker.xlsx:Recent Incidents row {source.source_row}","Safety_Manual.mbox:msg_safety_manual_5126","Workflow_Context.json:meta.current_time"]})
    return {"task_id":"automationbench-5126","status":"CONTROLLED_DRAFT_ONLY","frozen_time":FIXED_ISO_TIME,"position_count":8,"status_counts":{"Documented":3,"Under Review":2,"Investigation":3},"osha_reportable_count":3,"workers_comp_required_count":5,"email_draft_count":8,"emails_sent":0,"tracker_update_draft_count":8,"tracker_updates_applied":0,"osha_reports_filed":0,"workers_comp_forms_created":0,"slack_posts":0,"updates":entries}


def slack_text() -> str:
    return "Safety incident summary: 8 reviewed — 3 minor, 2 moderate, 3 serious; 3 require immediate Legal/OSHA review; 5 require workers' compensation follow-up. No employee names, manager addresses, medical details, or incident descriptions are included."


def slack_payload() -> dict:
    return {"task_id":"automationbench-5126","source_channel":{"id":"C_SAF_5126","name":"safety","is_private":False},"decision":"PUBLIC_AGGREGATE_DRAFT_ONLY","draft_count":1,"posted_count":0,"posted":False,"state":"DRAFT_NOT_POSTED","aggregate_counts":{"incidents":8,"minor":3,"moderate":2,"serious":3,"osha_legal_immediate":3,"workers_comp_follow_up":5},"contains_employee_names":False,"contains_manager_addresses":False,"contains_medical_details":False,"contains_incident_descriptions":False,"draft_text":slack_text(),"messages_created":[],"source_trace":["Slack_Channels.json:C_SAF_5126","Safety_Incident_Tracker.xlsx:Recent Incidents rows 2-9","Safety_Manual.mbox:msg_safety_manual_5126"]}


def execution_payload() -> dict:
    return {"task_id":"automationbench-5126","status":"CONTROLLED_DRAFT_ONLY","email":{"draft_count":8,"sent_count":0,"state":"DRAFT_NOT_SENT"},"tracker":{"update_draft_count":8,"updates_applied":0,"source_modified":False,"state":"UNCHANGED"},"slack":{"draft_count":1,"posted_count":0,"state":"DRAFT_NOT_POSTED"},"osha":{"reportable_count":3,"legal_notification_drafts":3,"reports_filed":0,"exact_deadlines_asserted":0,"reason":"Incident dates are available, but incident times are not."},"workers_comp":{"follow_up_required_count":5,"forms_created":0,"owner":"Benefits team separately"},"retention_years":5,"real_side_effect_count":0,"source_trace":["Safety_Incident_Tracker.xlsx:Recent Incidents rows 2-9","Safety_Manual.mbox:msg_safety_manual_5126","Workflow_Context.json:meta.current_time","Slack_Channels.json:C_SAF_5126"]}


def write_json(path: Path, payload: dict) -> None:
    path.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")


def build(input_dir: Path, output_dir: Path) -> None:
    records,manual,channels,frozen_time=read_inputs(input_dir); decisions=derive_decisions(records); output_dir.mkdir(parents=True,exist_ok=True)
    for child in output_dir.iterdir():
        if child.is_file(): child.unlink()
    write_workbook(output_dir/XLSX_NAME,records,manual,channels,frozen_time)
    write_brief(output_dir/DOCX_NAME,decisions)
    write_presentation(output_dir/PPTX_NAME,decisions)
    write_mbox(output_dir/MBOX_NAME,decisions)
    write_json(output_dir/UPDATES_NAME,updates_payload(decisions))
    write_json(output_dir/SLACK_NAME,slack_payload())
    write_json(output_dir/EXECUTION_NAME,execution_payload())


def main() -> None:
    parser=argparse.ArgumentParser(); parser.add_argument("--input-dir",type=Path,default=Path("/app/input_files")); parser.add_argument("--output-dir",type=Path,default=Path("/app/output")); args=parser.parse_args(); build(args.input_dir,args.output_dir)


if __name__ == "__main__":
    main()
