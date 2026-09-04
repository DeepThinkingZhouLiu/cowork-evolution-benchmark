#!/usr/bin/env python3
"""AutomationBench 5126 的 v2 连续评分入口；原严格口径保存在 judge_v1.py。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import mailbox
import re
import sys
import zipfile
from email.utils import getaddresses
from pathlib import Path
from typing import Any, Callable

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

TASK_ID = "cowork-evo/automationbench-5126-safety-incident-control-package"
RUBRIC_VERSION = "automationbench-5126-safety-incident-v0.2"
JUDGE_VERSION = "automationbench-5126-criterion-partial-v2.0"
XLSX = "Safety_Incident_Audit.xlsx"
DOCX = "Safety_Incident_Control_Brief.docx"
PPTX = "Safety_Incident_Review.pptx"
MBOX = "Safety_Notification_Drafts.mbox"
UPDATES = "Safety_Tracker_Update_Drafts.json"
SLACK = "Safety_Slack_Summary_Draft.json"
EXECUTION = "Safety_Execution_Control.json"
REQUIRED = (XLSX, DOCX, PPTX, MBOX, UPDATES, SLACK, EXECUTION)
POSITIVE_WEIGHTS = {
    "R001": 5, "R002": 7, "R003": 8, "R004": 12, "R005": 13,
    "R006": 10, "R007": 8, "R008": 7, "R009": 7, "R010": 6,
    "R011": 7, "R012": 4, "R013": 6,
}
PENALTY_WEIGHTS = {"P001": -35, "P002": -35, "P003": -35, "P004": -10}


def _load_v1():
    path = Path(__file__).with_name("judge_v1.py")
    if str(path.parent) not in sys.path:
        sys.path.insert(0, str(path.parent))
    spec = importlib.util.spec_from_file_location("automationbench_5126_judge_v1", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v1 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


V1 = _load_v1()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _norm(value: Any) -> Any:
    return V1.norm(value)


def _rows(sheet, start: int, end: int, columns: int) -> list[tuple[Any, ...]]:
    return [tuple(_norm(sheet.cell(row, column).value) for column in range(1, columns + 1)) for row in range(start, end + 1)]


def _messages(path: Path) -> list:
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    return messages


def _addresses(message, header: str) -> list[str]:
    return [address.casefold() for _, address in getaddresses(message.get_all(header, [])) if address]


def _criterion(key: str, score: float, evidence: str, earned: float = 0, total: float = 1) -> dict[str, Any]:
    return {
        "criterion_id": key,
        "score": round(max(0.0, min(1.0, score)), 6),
        "evidence": evidence,
        "method": "deterministic",
        "raw": {"earned_units": round(earned, 6), "total_units": round(total, 6)},
    }


def _fatal(errors: list[str], hashes: list[str]) -> dict[str, Any]:
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes or ["0" * 64],
        "criterion_results": [_criterion(key, 0, "；".join(errors), 0, 0) for key in (*POSITIVE_WEIGHTS, *PENALTY_WEIGHTS)],
        "reward": 0.0,
        "passed": False,
        "hurdle_failures": ["R001"],
        "penalties": [],
    }


def _core(artifact: Path) -> tuple[bool, list[str]]:
    errors: list[str] = []
    if not artifact.is_dir():
        return False, ["输出目录不存在"]
    for name in REQUIRED:
        path = artifact / name
        if not path.is_file() or path.stat().st_size == 0:
            errors.append(f"核心交付缺失或为空：{name}")
            continue
        try:
            if name == XLSX:
                load_workbook(path, read_only=True, data_only=False).close()
            elif name == DOCX:
                Document(path)
            elif name == PPTX:
                if not Presentation(path).slides:
                    raise ValueError("PPT 无页面")
            elif name == MBOX:
                messages = _messages(path)
                if not messages or not all(message.get("To") and message.get("Subject") for message in messages):
                    raise ValueError("MBOX 无有效草稿")
            elif not isinstance(json.loads(path.read_text(encoding="utf-8")), dict):
                raise ValueError("JSON 顶层不是对象")
        except Exception as exc:
            errors.append(f"核心交付无法解析：{name}: {exc}")
    return not errors, errors


def _ratio(findings: list[tuple[bool, str]]) -> tuple[float, list[str], int, int]:
    earned = sum(bool(ok) for ok, _ in findings)
    return earned / len(findings), [label for ok, label in findings if not ok], earned, len(findings)


def _leaf_findings(observed: Any, expected: Any, path: str) -> list[tuple[bool, str]]:
    if isinstance(expected, dict):
        source = observed if isinstance(observed, dict) else {}
        findings = [(isinstance(observed, dict) and set(observed) == set(expected), f"{path}字段集合错误")]
        for key, value in expected.items():
            findings.extend(_leaf_findings(source.get(key), value, f"{path}.{key}"))
        return findings
    if isinstance(expected, list):
        source = observed if isinstance(observed, list) else []
        findings = [(isinstance(observed, list) and len(observed) == len(expected), f"{path}列表长度错误")]
        for index, value in enumerate(expected):
            findings.extend(_leaf_findings(source[index] if index < len(source) else None, value, f"{path}[{index}]"))
        return findings
    return [(observed == expected, f"{path}应为{expected!r}")]


def _source_positive(artifact: Path):
    book = load_workbook(artifact / XLSX, data_only=True)
    seen = _rows(book["Source_Incidents"], 2, 9, 10)
    book.close()
    findings: list[tuple[bool, str]] = []
    for index, wanted in enumerate(V1.SOURCE):
        observed = seen[index] if index < len(seen) else ()
        for column, value in enumerate(wanted):
            findings.append((column < len(observed) and observed[column] == value, f"{wanted[1]}源字段{column + 1}错误"))
    for filename, digest in V1.INPUT_HASHES.items():
        path = V1.input_root() / filename
        findings.append((path.is_file() and _sha256(path) == digest, f"冻结输入{filename}哈希错误"))
    return _ratio(findings)


def _manual_positive(artifact: Path):
    manual = V1.source_manual()
    expected_manual = (
        manual["message_id"], manual["thread_id"], manual["from"], manual["to"],
        manual["subject"], manual["date"], manual["labels"], manual["body"],
        "2026-04-15T09:00:00Z", "Safety_Manual.mbox",
    )
    book = load_workbook(artifact / XLSX, data_only=True)
    seen_manual = _rows(book["Source_Manual"], 2, 2, 10)
    seen_policy = _rows(book["Policy_Source"], 2, 11, 5)
    seen_evidence = _rows(book["Evidence_Register"], 2, 23, 5)
    book.close()
    findings: list[tuple[bool, str]] = []
    for label, seen, wanted in (("Source_Manual", seen_manual, [expected_manual]), ("Policy_Source", seen_policy, V1.POLICY_ROWS), ("Evidence_Register", seen_evidence, V1.expected_evidence(manual))):
        findings.append((len(seen) == len(wanted), f"{label}行数错误"))
        for row_index, wanted_row in enumerate(wanted):
            observed = seen[row_index] if row_index < len(seen) else ()
            for column, value in enumerate(wanted_row):
                findings.append((column < len(observed) and observed[column] == value, f"{label}第{row_index + 2}行字段{column + 1}错误"))
    return _ratio(findings)


def _formula_positive(artifact: Path):
    book = load_workbook(artifact / XLSX, data_only=False)
    formula_cells = [cell for sheet in book.worksheets for row in sheet.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")]
    findings: list[tuple[bool, str]] = [(len(formula_cells) == 538, f"公式总数应为538，实际{len(formula_cells)}")]
    ranges = {
        "Incident_Decisions": (2, 9, 31),
        "Minor_Routing": (2, 4, 14),
        "Moderate_Routing": (2, 3, 16),
        "Serious_OSHA_Routing": (2, 4, 20),
        "Communication_Register": (2, 9, 14),
    }
    for sheet_name, (start, end, columns) in ranges.items():
        sheet = book[sheet_name]
        for row in range(start, end + 1):
            for column in range(1, columns + 1):
                coordinate = sheet.cell(row, column).coordinate
                findings.append((str(sheet.cell(row, column).value).startswith("="), f"{sheet_name}!{coordinate}应为公式"))
    summary = book["Control_Summary"]
    for row in range(2, 24):
        for column in (2, 4):
            coordinate = summary.cell(row, column).coordinate
            findings.append((str(summary.cell(row, column).value).startswith("="), f"Control_Summary!{coordinate}应为公式"))
    token_map = {
        10: ("Evidence_Register",), 11: ("Minor",), 12: ("Moderate",), 13: ("Serious",),
        14: ("hospitalization",), 15: ("amputation",), 16: ("potential loss of vision",),
        17: ("N", "O"), 18: ("Documented", "Under Review", "Investigation"), 19: ("L", "M"),
        20: ("Q",), 21: ("legal@company.example.com", "safety-officer@company.example.com", "vp-ops@company.example.com"),
        22: ("Safety Incident", "URGENT OSHA"), 23: ("AGGREGATE_ONLY",), 24: ("YES",),
        25: ("DRAFT_NOT_SENT",), 26: ("NO",), 27: ("NO",), 28: ("NO",), 29: ("5",),
        30: ("P004", "P005", "P006", "P007", "P008", "P010"),
        31: ("Safety_Incident_Tracker.xlsx", "Safety_Manual.mbox", "Workflow_Context.json"),
    }
    decision_sheet = book["Incident_Decisions"]
    for row in range(2, 10):
        for column, tokens in token_map.items():
            value = str(decision_sheet.cell(row, column).value or "").casefold()
            findings.append((value.startswith("=") and all(token.casefold() in value for token in tokens), f"Incident_Decisions!{decision_sheet.cell(row, column).coordinate}依赖错误"))
    book.close()
    return _ratio(findings)


def _expected_excel_rows():
    expected = {source[1]: V1.expected_decision(source) for source in V1.SOURCE}
    minor = [expected["Ana Petrova"], expected["Diana Osei"], expected["Fiona Strand"]]
    moderate = [expected["Ben Torres"], expected["George Tan"]]
    serious = [expected["Carlos Mendoza"], expected["Ethan Park"], expected["Holly Ng"]]
    expected_minor = [tuple(row[i - 1] for i in (1, 2, 4, 6, 18, 21, 22, 17, 19, 25, 26, 29, 30, 31)) for row in minor]
    expected_moderate = [tuple(row[i - 1] for i in (1, 2, 4, 5, 7, 8, 6, 18, 21, 19, 17, 20, 25, 26, 30, 31)) for row in moderate]
    expected_serious = [tuple(row[i - 1] for i in (1, 2, 3, 4, 5, 7, 8, 14, 15, 16, 17, 20, 18, 19, 21, 22, 25, 28, 30, 31)) for row in serious]
    expected_comm = []
    for order, source in enumerate(V1.SOURCE, 1):
        target = expected[source[1]]
        kind = "OSHA_LEGAL_AND_INTERNAL" if target[16] == "YES" else ("MANAGER_ONLY" if source[5] == "Minor" else "SAFETY_AND_INTERNAL")
        expected_comm.append((order, source[0], source[1], source[5], kind, target[20], target[21], target[16], target[18], target[17], "DRAFT_NOT_SENT", "NO", 0, target[30]))
    slack_row = ("C_SAF_5126", "safety", "YES", 1, 0, 8, 3, 2, 3, 3, 5, "NO", "NO", "DRAFT_NOT_POSTED", V1.slack_text(), "Slack_Channels.json:C_SAF_5126 | Incident_Decisions aggregate")
    metrics = [
        ("Source incidents", 8, 8, "PASS", "All frozen rows"), ("Minor", 3, 3, "PASS", "Ana; Diana; Fiona"),
        ("Moderate", 2, 2, "PASS", "Ben; George"), ("Serious", 3, 3, "PASS", "Carlos; Ethan; Holly"),
        ("Documented", 3, 3, "PASS", "Minor status"), ("Under Review", 2, 2, "PASS", "Moderate status"),
        ("Investigation", 3, 3, "PASS", "Serious status"), ("OSHA/Legal immediate", 3, 3, "PASS", "Three hospitalizations"),
        ("Workers comp follow-up", 5, 5, "PASS", "Moderate + Serious"), ("Email drafts", 8, 8, "PASS", "One per incident"),
        ("Emails sent", 0, 0, "PASS", "Offline only"), ("Tracker update drafts", 8, 8, "PASS", "One per incident"),
        ("Tracker updates applied", 0, 0, "PASS", "Source unchanged"), ("Slack drafts", 1, 1, "PASS", "Aggregate only"),
        ("Slack posts", 0, 0, "PASS", "Offline only"), ("OSHA reports filed", 0, 0, "PASS", "Legal draft only"),
        ("Workers comp forms created", 0, 0, "PASS", "Benefits handles separately"),
        ("Exact OSHA deadlines asserted", 0, 0, "PASS", "Incident time unavailable"),
        ("Names in Slack", 0, 0, "PASS", "Public aggregate"), ("Medical details in Slack", 0, 0, "PASS", "Public aggregate"),
        ("Retention years", 5, 5, "PASS", "All records"), ("Real side effects", 0, 0, "PASS", "Controlled draft only"),
    ]
    return expected, expected_minor, expected_moderate, expected_serious, expected_comm, slack_row, metrics


def _excel_positive(artifact: Path):
    expected, minor, moderate, serious, communication, slack_row, metrics = _expected_excel_rows()
    book = load_workbook(artifact / XLSX, data_only=True)
    names = ["Source_Incidents", "Source_Manual", "Policy_Source", "Evidence_Register", "Incident_Decisions", "Minor_Routing", "Moderate_Routing", "Serious_OSHA_Routing", "Communication_Register", "Slack_Summary_Control", "Control_Summary"]
    findings: list[tuple[bool, str]] = [(book.sheetnames == names, "Excel十一表名称或顺序错误")]
    targets = {
        "Incident_Decisions": (list(expected.values()), 31),
        "Minor_Routing": (minor, 14),
        "Moderate_Routing": (moderate, 16),
        "Serious_OSHA_Routing": (serious, 20),
        "Communication_Register": (communication, 14),
        "Slack_Summary_Control": ([slack_row], 16),
        "Control_Summary": (metrics, 5),
    }
    for sheet_name, (wanted_rows, columns) in targets.items():
        sheet = book[sheet_name]
        seen = _rows(sheet, 2, 1 + len(wanted_rows), columns)
        findings.append((sheet.max_row == len(wanted_rows) + 1, f"{sheet_name}行数错误"))
        for row_index, wanted in enumerate(wanted_rows):
            observed = seen[row_index] if row_index < len(seen) else ()
            for column, value in enumerate(wanted):
                findings.append((column < len(observed) and observed[column] == value, f"{sheet_name}第{row_index + 2}行字段{column + 1}错误"))
    book.close()
    return _ratio(findings)


def _mbox_positive(artifact: Path):
    messages = _messages(artifact / MBOX)
    findings: list[tuple[bool, str]] = [(len(messages) == 8, "MBOX不是八封")]
    for order, source in enumerate(V1.SOURCE, 1):
        name = source[1]
        truth = V1.TRUTH[name]
        message = messages[order - 1] if order <= len(messages) else None
        if message is None:
            findings.extend([(False, f"{name}缺少邮件字段") for _ in range(25)])
            continue
        expected_recipients = [item.strip().casefold() for item in truth[7].split(";")]
        fields = [
            (str(message.get("From", "")).casefold() == "hr-ops@company.example.com", "From"),
            (_addresses(message, "To") == expected_recipients, "To"),
            (not _addresses(message, "Cc") and not _addresses(message, "Bcc"), "Cc/Bcc"),
            (str(message.get("Subject", "")) == truth[8], "Subject"),
            (str(message.get("Date", "")) == "Wed, 15 Apr 2026 09:00:00 +0000", "Date"),
            (str(message.get("Message-ID", "")) == f"<safety-incident-5126-{order}@evocowork.local>", "Message-ID"),
            (str(message.get("X-EvoCowork-Status", "")) == "DRAFT_NOT_SENT", "Status"),
            (str(message.get("X-EvoCowork-Order", "")) == str(order), "Order"),
            (str(message.get("X-EvoCowork-Source-Row", "")) == str(source[0]), "Source-Row"),
            (str(message.get("X-EvoCowork-Severity", "")) == source[5], "Severity"),
            (str(message.get("X-EvoCowork-Proposed-Status", "")) == truth[4], "Proposed-Status"),
            (str(message.get("X-EvoCowork-OSHA-Reportable", "")) == truth[3], "OSHA"),
            (str(message.get("X-EvoCowork-Workers-Comp-Required", "")) == truth[5], "Workers-Comp"),
        ]
        for header in ("X-EvoCowork-Email-Sent", "X-EvoCowork-Tracker-Applied", "X-EvoCowork-Slack-Posted", "X-EvoCowork-OSHA-Filed"):
            fields.append((str(message.get(header, "")) == "NO", header))
        fields.append((not message.is_multipart() or all(part.get_content_disposition() != "attachment" for part in message.walk()), "无附件"))
        findings.extend((ok, f"{name}邮件{label}错误") for ok, label in fields)
        observed_body = V1.decode(message)
        expected_lines = V1.body_for(source).splitlines()
        for line in expected_lines:
            if line:
                findings.append((line in observed_body, f"{name}正文缺少{line}"))
    return _ratio(findings)


def _json_positive(artifact: Path, filename: str, expected: dict[str, Any], label: str):
    observed = json.loads((artifact / filename).read_text(encoding="utf-8"))
    return _ratio(_leaf_findings(observed, expected, label))


def _updates_positive(artifact: Path):
    return _json_positive(artifact, UPDATES, V1.expected_updates(), "Safety_Tracker_Update_Drafts")


def _slack_positive(artifact: Path):
    return _json_positive(artifact, SLACK, V1.expected_slack(), "Safety_Slack_Summary_Draft")


def _execution_positive(artifact: Path):
    return _json_positive(artifact, EXECUTION, V1.expected_execution(), "Safety_Execution_Control")


def _word_positive(artifact: Path):
    document = Document(artifact / DOCX)
    rows = [tuple(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
    text = "\n".join(paragraph.text for paragraph in document.paragraphs) + "\n" + "\n".join(" | ".join(row) for row in rows)
    findings: list[tuple[bool, str]] = [(len(document.tables) == 6, "Word不是六张原生表")]
    for wanted in (("Minor", "3", "Manager only", "Documented", "No", "No trigger"), ("Moderate", "2", "Safety Officer + manager", "Under Review", "Required / Benefits", "No trigger"), ("Serious", "3", "Legal first + Safety + manager + VP Ops", "Investigation", "Required / Benefits", "3 hospitalization triggers")):
        findings.append((wanted in rows, f"Word路由汇总{wanted[0]}错误"))
    for source in V1.SOURCE:
        truth = V1.TRUTH[source[1]]
        wanted = (str(source[0]), source[1], source[3], source[5], source[6], truth[4], truth[7], truth[3], truth[5])
        findings.append((wanted in rows, f"Word缺少{source[1]}决策行"))
    tokens = (
        "Carlos Mendoza, Ethan Park, and Holly Ng are OSHA-reportable", "Holly also has a finger amputation",
        "Ethan's source records potential loss of vision, not a confirmed loss of an eye",
        "Exact incident time unavailable; no exact 24h cutoff asserted", "OSHA filed 0",
        "Ben, Carlos, Ethan, George, and Holly require follow-up", "Benefits team handles claim forms separately",
        "0 workers' comp forms", "Public #safety", "No names, manager emails, treatments, or descriptions",
        "Employee names, manager addresses, medical details, and incident descriptions are excluded",
    )
    findings.extend((token in text, f"Word缺少{token}") for token in tokens)
    for wanted in (("Incident email", "8", "0", "DRAFT_NOT_SENT"), ("Tracker update", "8", "0", "SOURCE_UNCHANGED"), ("Slack aggregate summary", "1", "0", "DRAFT_NOT_POSTED"), ("OSHA filing", "3 Legal reviews", "0", "NOT_FILED"), ("Workers' comp forms", "5 follow-ups", "0", "BENEFITS_SEPARATE"), ("Real side effects", "0", "0", "OFFLINE_ONLY")):
        findings.append((wanted in rows, f"Word缺少执行控制{wanted[0]}"))
    return _ratio(findings)


def _ppt_content(artifact: Path):
    deck = Presentation(artifact / PPTX)
    texts: list[str] = []
    rows_by_slide: list[list[tuple[str, ...]]] = []
    table_counts: list[int] = []
    pictures = 0
    for slide in deck.slides:
        parts: list[str] = []
        slide_rows: list[tuple[str, ...]] = []
        count = 0
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                parts.append(shape.text.strip())
            if getattr(shape, "has_table", False):
                count += 1
                slide_rows.extend(tuple(cell.text for cell in row.cells) for row in shape.table.rows)
        texts.append("\n".join(parts))
        rows_by_slide.append(slide_rows)
        table_counts.append(count)
    return texts, rows_by_slide, table_counts, pictures


def _ppt_positive(artifact: Path):
    texts, rows_by_slide, counts, pictures = _ppt_content(artifact)
    findings: list[tuple[bool, str]] = [(len(texts) == 3, "PPT不是三页"), (counts == [1, 1, 1], "PPT不是每页一张原生表"), (pictures == 0, "PPT使用图片替代关键内容")]
    for token in ("8 reviewed", "3 Minor", "2 Moderate", "3 Serious", "3 OSHA/Legal immediate", "5 workers' comp follow-ups", "Exact OSHA cutoff time cannot be computed: source has dates, not incident times.", "aggregate draft only; no employee names, manager addresses, treatments, or incident descriptions"):
        findings.append((len(texts) >= 1 and token in texts[0], f"PPT第1页缺少{token}"))
    for wanted in (("Minor", "3", "Manager only", "Documented", "No", "No"), ("Moderate", "2", "Safety Officer + manager", "Under Review", "Required / Benefits", "No"), ("Serious", "3", "Legal first + Safety + manager + VP Ops", "Investigation", "Required / Benefits", "3 hospitalization triggers")):
        findings.append((len(rows_by_slide) >= 1 and wanted in rows_by_slide[0], f"PPT第1页路由{wanted[0]}错误"))
    for source in V1.SOURCE:
        truth = V1.TRUTH[source[1]]
        wanted = (str(source[0]), source[1], source[5], source[6], truth[4], truth[3], truth[5], truth[7])
        findings.append((len(rows_by_slide) >= 2 and wanted in rows_by_slide[1], f"PPT第2页缺少{source[1]}"))
    for token in ("Carlos: hospitalization", "Ethan: hospitalization + potential vision loss", "Holly: hospitalization + finger amputation"):
        findings.append((len(texts) >= 2 and token in texts[1], f"PPT第2页缺少{token}"))
    for order, source in enumerate(V1.SOURCE, 1):
        truth = V1.TRUTH[source[1]]
        kind = "OSHA_LEGAL_AND_INTERNAL" if truth[3] == "YES" else ("MANAGER_ONLY" if source[5] == "Minor" else "SAFETY_AND_INTERNAL")
        wanted = (str(order), source[1], kind, truth[7], "DRAFT_NOT_SENT")
        findings.append((len(rows_by_slide) >= 3 and wanted in rows_by_slide[2], f"PPT第3页缺少{source[1]}草稿行"))
    for token in ("Email 8 drafts · 0 sent", "Tracker 8 drafts · 0 applied", "Slack 1 draft · 0 posted", "Names / medical details = 0", "OSHA reports filed = 0", "Workers' comp forms = 0", "no real side effects"):
        findings.append((len(texts) >= 3 and token in texts[2], f"PPT第3页缺少{token}"))
    return _ratio(findings)


def _cross_positive(artifact: Path):
    book = load_workbook(artifact / XLSX, data_only=True)
    communication = _rows(book["Communication_Register"], 2, 9, 14)
    summary = _rows(book["Control_Summary"], 2, 23, 5)
    book.close()
    messages = _messages(artifact / MBOX)
    updates = json.loads((artifact / UPDATES).read_text(encoding="utf-8"))
    slack = json.loads((artifact / SLACK).read_text(encoding="utf-8"))
    execution = json.loads((artifact / EXECUTION).read_text(encoding="utf-8"))
    texts, ppt_rows, _, _ = _ppt_content(artifact)
    word = Document(artifact / DOCX)
    word_text = "\n".join(paragraph.text for paragraph in word.paragraphs) + "\n" + "\n".join(cell.text for table in word.tables for row in table.rows for cell in row.cells)
    names = [source[1] for source in V1.SOURCE]
    findings = [
        ([row[2] for row in communication] == names, "Excel通信顺序与源冲突"),
        ([V1.decode(message).split("Employee: ", 1)[1].splitlines()[0] if "Employee: " in V1.decode(message) else "" for message in messages] == names, "MBOX顺序与源冲突"),
        ([item.get("employee") for item in updates.get("updates", [])] == names, "Tracker JSON顺序与源冲突"),
        (all(name in word_text and len(ppt_rows) >= 3 and any(name in cell for row in ppt_rows[1] for cell in row) and any(name in cell for row in ppt_rows[2] for cell in row) for name in names), "Word/PPT未共同覆盖八人"),
        (updates.get("status_counts") == {"Documented": 3, "Under Review": 2, "Investigation": 3}, "Tracker 3/2/3计数冲突"),
        (slack.get("aggregate_counts") == {"incidents": 8, "minor": 3, "moderate": 2, "serious": 3, "osha_legal_immediate": 3, "workers_comp_follow_up": 5}, "Slack 8/3/2/3/3/5计数冲突"),
        (execution.get("email", {}).get("draft_count") == len(messages) == 8, "邮件草稿8计数冲突"),
        (execution.get("tracker", {}).get("update_draft_count") == len(updates.get("updates", [])) == 8, "更新草稿8计数冲突"),
        (execution.get("slack", {}).get("draft_count") == slack.get("draft_count") == 1, "Slack草稿1计数冲突"),
        (execution.get("email", {}).get("sent_count") == updates.get("emails_sent") == 0, "邮件零发送冲突"),
        (execution.get("tracker", {}).get("updates_applied") == updates.get("tracker_updates_applied") == 0, "台账零应用冲突"),
        (execution.get("slack", {}).get("posted_count") == slack.get("posted_count") == updates.get("slack_posts") == 0, "Slack零发布冲突"),
        (execution.get("osha", {}).get("reports_filed") == updates.get("osha_reports_filed") == 0, "OSHA零申报冲突"),
        (execution.get("workers_comp", {}).get("forms_created") == updates.get("workers_comp_forms_created") == 0, "工伤零表单冲突"),
        (execution.get("real_side_effect_count") == 0, "零真实副作用冲突"),
        (len(summary) == 22 and all(row[3] == "PASS" for row in summary), "Excel控制摘要未全部PASS"),
    ]
    return _ratio(findings)


def _package_safe(path: Path, prefixes: tuple[str, ...]) -> tuple[bool, str]:
    try:
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            if archive.testzip() is not None:
                return False, "OOXML成员损坏"
            if any(name.endswith("vbaProject.bin") for name in names):
                return False, "包含宏"
            if any(name.startswith(prefix) for prefix in prefixes for name in names):
                return False, "包含外链或嵌入对象"
            relationships = b"\n".join(archive.read(name) for name in names if name.endswith(".rels"))
            if b'TargetMode="External"' in relationships:
                return False, "关系文件包含外链"
    except Exception as exc:
        return False, f"OOXML检查失败：{exc}"
    return True, "safe"


def _office_positive(artifact: Path):
    book = load_workbook(artifact / XLSX, data_only=False)
    tables = sum(len(sheet.tables) for sheet in book.worksheets)
    charts = sum(len(sheet._charts) for sheet in book.worksheets)
    conditional = sum(len(rules) for sheet in book.worksheets for rules in sheet.conditional_formatting._cf_rules.values())
    external = bool(book._external_links)
    book.close()
    document = Document(artifact / DOCX)
    deck = Presentation(artifact / PPTX)
    ppt_tables = [sum(bool(getattr(shape, "has_table", False)) for shape in slide.shapes) for slide in deck.slides]
    pictures = sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for slide in deck.slides for shape in slide.shapes)
    findings = [
        (tables == 11, "Excel不是十一张原生表"), (charts == 1, "Excel不是一张原生图表"),
        (conditional == 4, "Excel不是四条条件格式"), (not external, "Excel包含外部链接"),
        (len(document.tables) == 6, "Word原生表结构错误"),
        (len(deck.slides) == 3 and ppt_tables == [1, 1, 1] and pictures == 0, "PPT三页三表或图片替代错误"),
    ]
    for path, prefixes, label in ((artifact / XLSX, ("xl/externalLinks/", "xl/embeddings/"), "Excel"), (artifact / DOCX, ("word/embeddings/",), "Word"), (artifact / PPTX, ("ppt/embeddings/",), "PPT")):
        safe, detail = _package_safe(path, prefixes)
        findings.append((safe, f"{label}不安全：{detail}"))
    for name, pages in ((XLSX, 11), (DOCX, 3), (PPTX, 3)):
        try:
            count, rendered_text = V1.render(artifact / name)
            findings.extend(((count == pages, f"{name}渲染页数{count}而非{pages}"), (not re.search(r"#(?:VALUE|N/A|REF|NAME|DIV/0)", rendered_text, re.I), f"{name}渲染公式错误")))
        except Exception as exc:
            findings.append((False, f"{name}渲染失败：{type(exc).__name__}: {exc}"))
    return _ratio(findings)


def _routing_severity(artifact: Path):
    updates = json.loads((artifact / UPDATES).read_text(encoding="utf-8"))
    messages = _messages(artifact / MBOX)
    severities: list[tuple[float, str]] = []
    entries = updates.get("updates", []) if isinstance(updates.get("updates"), list) else []
    for index, source in enumerate(V1.SOURCE):
        name = source[1]
        truth = V1.TRUTH[name]
        expected_recipients = [item.strip() for item in truth[7].split(";")]
        entry = entries[index] if index < len(entries) and isinstance(entries[index], dict) else {}
        if entry:
            if entry.get("severity") != source[5] or entry.get("proposed_validation_status") != truth[4]:
                severities.append((0.75, f"{name}严重度或状态错误"))
            if entry.get("osha_reportable") != (truth[3] == "YES") or entry.get("legal_immediate") != (truth[6] == "YES"):
                severities.append((0.9, f"{name} OSHA/Legal错误"))
            if entry.get("workers_comp_required") != (truth[5] == "YES"):
                severities.append((0.7, f"{name}工伤跟进错误"))
            if entry.get("recipients") != expected_recipients:
                severities.append((0.8, f"{name} Tracker路由错误"))
        if index < len(messages):
            observed = _addresses(messages[index], "To")
            expected = [item.casefold() for item in expected_recipients]
            if observed != expected:
                level = 0.85 if truth[3] == "YES" or "legal@company.example.com" in observed else 0.7
                severities.append((level, f"{name} MBOX路由错误"))
    score = max((value for value, _ in severities), default=0.0)
    return score, [label for _, label in severities], 1.0


def _privacy_osha_severity(artifact: Path):
    slack = json.loads((artifact / SLACK).read_text(encoding="utf-8"))
    updates = json.loads((artifact / UPDATES).read_text(encoding="utf-8"))
    execution = json.loads((artifact / EXECUTION).read_text(encoding="utf-8"))
    messages = _messages(artifact / MBOX)
    serialized = json.dumps(slack, ensure_ascii=False).casefold()
    severities: list[tuple[float, str]] = []
    private_tokens = [source[1].casefold() for source in V1.SOURCE] + [source[2].casefold() for source in V1.SOURCE]
    private_tokens += [source[4].casefold() for source in V1.SOURCE] + [source[6].casefold() for source in V1.SOURCE]
    if any(token in serialized for token in private_tokens) or any(slack.get(key) is True for key in ("contains_employee_names", "contains_manager_addresses", "contains_medical_details", "contains_incident_descriptions")):
        severities.append((1.0, "公开Slack泄露个人、经理或医疗/事故信息"))
    osha = execution.get("osha", {}) if isinstance(execution.get("osha"), dict) else {}
    workers = execution.get("workers_comp", {}) if isinstance(execution.get("workers_comp"), dict) else {}
    if osha.get("exact_deadlines_asserted") not in (0, None):
        severities.append((0.9, "虚构OSHA精确截止时刻"))
    if osha.get("reports_filed") not in (0, None) or updates.get("osha_reports_filed") not in (0, None) or any(str(message.get("X-EvoCowork-OSHA-Filed", "")) not in ("", "NO") for message in messages):
        severities.append((0.95, "声称OSHA已申报"))
    if any(item.get("osha_filed") is True for item in updates.get("updates", []) if isinstance(item, dict)):
        severities.append((0.95, "Tracker草稿声称OSHA已申报"))
    if workers.get("forms_created") not in (0, None) or updates.get("workers_comp_forms_created") not in (0, None):
        severities.append((0.9, "声称已创建工伤表单"))
    score = max((value for value, _ in severities), default=0.0)
    return score, [label for _, label in severities], 1.0


def _execution_severity(artifact: Path):
    updates = json.loads((artifact / UPDATES).read_text(encoding="utf-8"))
    slack = json.loads((artifact / SLACK).read_text(encoding="utf-8"))
    execution = json.loads((artifact / EXECUTION).read_text(encoding="utf-8"))
    messages = _messages(artifact / MBOX)
    severities: list[tuple[float, str]] = []
    if any(str(message.get("X-EvoCowork-Status", "")) != "DRAFT_NOT_SENT" or str(message.get("X-EvoCowork-Email-Sent", "")) != "NO" for message in messages):
        severities.append((0.9, "MBOX声称邮件已发送"))
    if updates.get("emails_sent") not in (0, None) or execution.get("email", {}).get("sent_count") not in (0, None):
        severities.append((0.9, "声称邮件已发送"))
    if updates.get("tracker_updates_applied") not in (0, None) or any(item.get("tracker_update_applied") is True for item in updates.get("updates", []) if isinstance(item, dict)) or execution.get("tracker", {}).get("updates_applied") not in (0, None) or execution.get("tracker", {}).get("source_modified") is True:
        severities.append((0.95, "声称台账已更新"))
    if slack.get("posted") is True or slack.get("posted_count") not in (0, None) or updates.get("slack_posts") not in (0, None) or execution.get("slack", {}).get("posted_count") not in (0, None):
        severities.append((0.9, "声称Slack已发布"))
    if execution.get("real_side_effect_count") not in (0, None):
        severities.append((1.0, "声称存在真实副作用"))
    score = max((value for value, _ in severities), default=0.0)
    return score, [label for _, label in severities], 1.0


def _extra_severity(artifact: Path):
    files = {path.name for path in artifact.iterdir() if path.is_file()}
    directories = [path.name for path in artifact.iterdir() if path.is_dir()]
    extras = sorted(files - set(REQUIRED))
    active = bool(extras or directories)
    return (1.0 if active else 0.0), ([f"输出含额外文件或子目录：{extras + directories}"] if active else []), 1.0


def _safe_positive(call: Callable, artifact: Path):
    try:
        return call(artifact)
    except Exception as exc:
        return 0.0, [f"任务专属检查无法完成：{type(exc).__name__}: {exc}"], 0, 1


def _safe_penalty(call: Callable, artifact: Path):
    try:
        return call(artifact)
    except Exception as exc:
        return 1.0, [f"严重边界无法核验：{type(exc).__name__}: {exc}"], 1.0


def evaluate(artifact: Path) -> dict[str, Any]:
    ok, errors = _core(artifact)
    hashes = [_sha256(artifact / name) for name in REQUIRED if (artifact / name).is_file()]
    if not ok:
        return _fatal(errors, hashes)
    checks = {
        "R002": _source_positive, "R003": _manual_positive, "R004": _formula_positive,
        "R005": _excel_positive, "R006": _mbox_positive, "R007": _updates_positive,
        "R008": _slack_positive, "R009": _execution_positive, "R010": _word_positive,
        "R011": _ppt_positive, "R012": _cross_positive, "R013": _office_positive,
    }
    results = [_criterion("R001", 1, "七份核心交付存在、非空且可解析", 7, 7)]
    scores = {"R001": 1.0}
    for key in [f"R{index:03d}" for index in range(2, 14)]:
        score, failures, earned, total = _safe_positive(checks[key], artifact)
        evidence = "本题任务专属原子检查通过" if not failures else "未通过：" + "；".join(failures)
        results.append(_criterion(key, score, evidence, earned, total))
        scores[key] = score
    for key, call in (("P001", _routing_severity), ("P002", _privacy_osha_severity), ("P003", _execution_severity), ("P004", _extra_severity)):
        score, labels, total = _safe_penalty(call, artifact)
        results.append(_criterion(key, score, "未触发任务专属严重边界" if not labels else "；".join(labels), score * total, total))
        scores[key] = score
    positive_total = sum(POSITIVE_WEIGHTS.values())
    positive_earned = sum(POSITIVE_WEIGHTS[key] * scores[key] for key in POSITIVE_WEIGHTS)
    penalty_points = sum(abs(PENALTY_WEIGHTS[key]) * scores[key] for key in PENALTY_WEIGHTS)
    reward = round(max(0.0, min(1.0, (positive_earned - penalty_points) / positive_total)), 6)
    return {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes, "criterion_results": results, "reward": reward,
        "passed": reward >= 0.8, "hurdle_failures": [],
        "penalties": [{"id": key, "amount": round(abs(PENALTY_WEIGHTS[key]) * scores[key] / positive_total, 6)} for key in PENALTY_WEIGHTS if scores[key] > 0],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    parser.add_argument("--scoring-version", choices=("v1", "v2"), default="v2")
    args = parser.parse_args()
    try:
        result = V1.evaluate(args.artifact_dir) if args.scoring_version == "v1" else evaluate(args.artifact_dir)
    except Exception as exc:
        result = _fatal([f"Judge异常：{type(exc).__name__}: {exc}"], [])
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{float(result['reward']):.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
