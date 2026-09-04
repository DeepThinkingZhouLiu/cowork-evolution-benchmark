#!/usr/bin/env python3
"""AutomationBench 5126 的任务专属 v2 分层负例。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from typing import Callable

from docx import Document
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution/Golden"
INPUTS = ROOT / "environment/input_files"
OUTPUTS = (
    "Safety_Incident_Audit.xlsx", "Safety_Incident_Control_Brief.docx",
    "Safety_Incident_Review.pptx", "Safety_Notification_Drafts.mbox",
    "Safety_Tracker_Update_Drafts.json", "Safety_Slack_Summary_Draft.json",
    "Safety_Execution_Control.json",
)
XLSX, DOCX, PPTX, MBOX, UPDATES, SLACK, EXECUTION = OUTPUTS
Mutator = Callable[[Path], None]


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"无法载入{path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


V1_NEG = _load("safety_5126_negative_v1_helpers", ROOT / "tests/negative_cases_v1.py")


def _input_only(root: Path) -> None:
    for path in list(root.iterdir()):
        if path.is_file():
            path.unlink()
    for source in INPUTS.iterdir():
        if source.is_file():
            shutil.copy2(source, root / source.name)


def _excel(sheet: str, cell: str, value=None, *, remove_formula: bool = False) -> Mutator:
    def mutate(root: Path) -> None:
        V1_NEG.edit_cell(root / XLSX, sheet, cell, value=value, remove_formula=remove_formula)
    return mutate


def _json(filename: str, mutator: Callable[[dict], None]) -> Mutator:
    def mutate(root: Path) -> None:
        path = root / filename
        payload = json.loads(path.read_text(encoding="utf-8"))
        mutator(payload)
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return mutate


def _message(index: int, mutator: Callable) -> Mutator:
    def mutate(root: Path) -> None:
        path = root / MBOX
        box = mailbox.mbox(path, create=False)
        messages = list(box)
        box.close()
        mutator(messages[index])
        with path.open("wb") as handle:
            for order, message in enumerate(messages, 1):
                handle.write(f"From MAILER-DAEMON Wed Apr 15 09:00:{order:02d} 2026\n".encode("ascii"))
                BytesGenerator(handle, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
                handle.write(b"\n")
    return mutate


def _replace_body(message, old: str, new: str) -> None:
    body = V1_NEG.decode_message(message).replace(old, new, 1)
    for header in ("Content-Type", "Content-Transfer-Encoding", "MIME-Version"):
        if header in message:
            del message[header]
    message.set_content(body.rstrip() + "\n", charset="utf-8")


def _word_time_boundary(root: Path) -> None:
    path = root / DOCX
    document = Document(path)
    old = "Exact incident time unavailable; no exact 24h cutoff asserted"
    changed = False
    for paragraph in document.paragraphs:
        if old in paragraph.text:
            paragraph.text = paragraph.text.replace(old, "Exact incident time unavailable")
            changed = True
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if old in cell.text:
                    cell.text = cell.text.replace(old, "Exact incident time unavailable")
                    changed = True
    if not changed:
        raise ValueError("Word中未找到OSHA时刻边界")
    document.save(path)


def _ppt_george_status(root: Path) -> None:
    path = root / PPTX
    deck = Presentation(path)
    changed = False
    for slide in deck.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_table", False):
                continue
            for row in shape.table.rows:
                if any(cell.text == "George Tan" for cell in row.cells):
                    for cell in row.cells:
                        if cell.text == "Under Review":
                            cell.text = "Documented"
                            changed = True
                            break
                if changed:
                    break
            if changed:
                break
        if changed:
            break
    if not changed:
        raise ValueError("PPT中未找到George状态")
    deck.save(path)


def _cases():
    cases = []
    for filename in OUTPUTS:
        cases.append((f"missing-{filename}", lambda root, item=filename: (root / item).unlink(), "fatal_zero"))
        cases.append((f"corrupt-{filename}", lambda root, item=filename: (root / item).write_bytes(b"not-a-valid-artifact"), "fatal_zero"))
    cases.extend([
        ("mbox-empty", lambda root: (root / MBOX).write_bytes(b""), "fatal_zero"),
        ("unmodified-input-only", _input_only, "fatal_zero"),
        ("source-ana-date-local", _excel("Source_Incidents", "D2", "2026-04-15"), "partial"),
        ("manual-subject-local", _excel("Source_Manual", "E2", "Workplace Safety Manual"), "partial"),
        ("formula-carlos-hospitalization-local", _excel("Incident_Decisions", "N4", remove_formula=True), "partial"),
        ("minor-ana-recipient-local", _excel("Minor_Routing", "F2", "safety-officer@company.example.com"), "partial"),
        ("moderate-ben-workers-comp-local", _excel("Moderate_Routing", "J2", "NO"), "partial"),
        ("serious-holly-amputation-local", _excel("Serious_OSHA_Routing", "I4", "NO"), "partial"),
        ("mbox-ana-subject-local", _message(0, lambda message: message.replace_header("Subject", "Safety Incident - Ana Petrova")), "partial"),
        ("mbox-carlos-legal-omitted", _message(2, lambda message: message.replace_header("To", "safety-officer@company.example.com, mgr-a@company.example.com, vp-ops@company.example.com")), "partial"),
        ("updates-ethan-vision-local", _json(UPDATES, lambda payload: payload["updates"][4].update({"vision_source": "CONFIRMED_EYE_LOSS"})), "partial"),
        ("updates-holly-osha-filed", _json(UPDATES, lambda payload: payload["updates"][7].update({"osha_filed": True})), "partial"),
        ("slack-minor-count-local", _json(SLACK, lambda payload: payload["aggregate_counts"].update({"minor": 4})), "partial"),
        ("slack-employee-name-leak", _json(SLACK, lambda payload: payload.update({"draft_text": payload["draft_text"] + " Ana Petrova."})), "partial"),
        ("execution-email-draft-count-local", _json(EXECUTION, lambda payload: payload["email"].update({"draft_count": 7})), "partial"),
        ("word-osha-time-boundary-local", _word_time_boundary, "partial"),
        ("ppt-george-status-local", _ppt_george_status, "partial"),
        ("extra-file", lambda root: (root / "unexpected.txt").write_text("unexpected", encoding="utf-8"), "partial"),
    ])
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    judge = _load("safety_5126_v2_judge", ROOT / "tests/judge.py")
    judge.V1.render = lambda path: ({".xlsx": 11, ".docx": 3, ".pptx": 3}[path.suffix.casefold()], "render-ok")
    if judge.evaluate(GOLDEN).get("reward") != 1.0:
        raise AssertionError("Golden v2必须为1.0")
    results = []
    for name, mutator, layer in _cases():
        with tempfile.TemporaryDirectory(prefix=f"safety-5126-v2-{name}-") as temporary:
            output = Path(temporary) / "output"
            shutil.copytree(GOLDEN, output)
            mutator(output)
            result = judge.evaluate(output)
            reward = float(result.get("reward", 0))
            if layer == "fatal_zero" and reward != 0:
                raise AssertionError(f"{name}应硬归零，实际{reward}")
            if layer == "partial" and not 0 < reward < 1:
                raise AssertionError(f"{name}应保留连续部分分，实际{reward}")
            results.append({"case": name, "layer": layer, "reward": reward, "passed": bool(result.get("passed")), "penalties": result.get("penalties", [])})
    partial = [item["reward"] for item in results if item["layer"] == "partial"]
    report = {
        "total": len(results), "fatal_zero": sum(item["layer"] == "fatal_zero" for item in results),
        "partial_continuous": len(partial), "partial_reward_range": [min(partial), max(partial)], "cases": results,
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
