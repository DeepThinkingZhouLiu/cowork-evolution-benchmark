#!/usr/bin/env python3
"""逐项注入 AutomationBench 5126 专属错误，验证 Judge 不会误接受坏答案。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
import zipfile
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path
from typing import Callable

from docx import Document
from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pptx import Presentation


ROOT=Path(__file__).resolve().parents[1]; GOLDEN=ROOT/"solution"/"Golden"
XLSX="Safety_Incident_Audit.xlsx"; DOCX="Safety_Incident_Control_Brief.docx"; PPTX="Safety_Incident_Review.pptx"; MBOX="Safety_Notification_Drafts.mbox"; UPDATES="Safety_Tracker_Update_Drafts.json"; SLACK="Safety_Slack_Summary_Draft.json"; EXECUTION="Safety_Execution_Control.json"
FILES=(XLSX,DOCX,PPTX,MBOX,UPDATES,SLACK,EXECUTION); Mutator=Callable[[Path],None]; Case=tuple[str,Mutator,str]


def load_judge():
    spec=importlib.util.spec_from_file_location("safety_5126_judge",ROOT/"tests"/"judge.py")
    if spec is None or spec.loader is None: raise RuntimeError("无法载入本题Judge")
    module=importlib.util.module_from_spec(spec); sys.modules[spec.name]=module; spec.loader.exec_module(module); module.render=lambda path:({".xlsx":11,".docx":3,".pptx":3}[path.suffix],"render-ok"); return module


def rewrite_zip(path:Path,transform:Callable[[str,bytes],bytes|None])->None:
    temporary=path.with_suffix(path.suffix+".negative")
    with zipfile.ZipFile(path,"r") as source,zipfile.ZipFile(temporary,"w",zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data=transform(info.filename,source.read(info.filename))
            if data is not None:target.writestr(info,data)
    temporary.replace(path)


def sheet_xml_path(archive:zipfile.ZipFile,sheet_name:str)->str:
    workbook=etree.fromstring(archive.read("xl/workbook.xml")); relationships=etree.fromstring(archive.read("xl/_rels/workbook.xml.rels")); ns={"m":"http://schemas.openxmlformats.org/spreadsheetml/2006/main","r":"http://schemas.openxmlformats.org/officeDocument/2006/relationships","pr":"http://schemas.openxmlformats.org/package/2006/relationships"}
    relationship_id=workbook.xpath(f'//m:sheet[@name="{sheet_name}"]/@r:id',namespaces=ns)[0]; target=relationships.xpath(f'//pr:Relationship[@Id="{relationship_id}"]/@Target',namespaces=ns)[0]; return target.lstrip("/") if target.startswith("/") else "xl/"+target.lstrip("/")


def edit_cell(path:Path,sheet_name:str,coordinate:str,*,value=None,remove_formula:bool=False)->None:
    with zipfile.ZipFile(path,"r") as archive:target_sheet=sheet_xml_path(archive,sheet_name)
    changed=False; namespace="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    def transform(name:str,data:bytes)->bytes:
        nonlocal changed
        if name!=target_sheet:return data
        root=etree.fromstring(data); cells=root.xpath(f'//m:c[@r="{coordinate}"]',namespaces={"m":namespace})
        if not cells:raise ValueError(f"{sheet_name}!{coordinate}不存在")
        cell=cells[0]
        if remove_formula:
            node=cell.find(f"{{{namespace}}}f")
            if node is None:raise ValueError(f"{sheet_name}!{coordinate}没有公式")
            cell.remove(node)
        if value is not None:
            inline=cell.find(f"{{{namespace}}}is")
            if inline is not None:cell.remove(inline)
            node=cell.find(f"{{{namespace}}}v")
            if node is None:node=etree.SubElement(cell,f"{{{namespace}}}v")
            if isinstance(value,bool):node.text="1" if value else "0";cell.set("t","b")
            elif isinstance(value,(int,float)):node.text=str(value);cell.attrib.pop("t",None)
            else:node.text=str(value);cell.set("t","str")
        changed=True;return etree.tostring(root,xml_declaration=True,encoding="UTF-8",standalone=True)
    rewrite_zip(path,transform)
    if not changed:raise ValueError("目标单元格未修改")


def openpyxl_mutation(path:Path,mutator:Callable)->None:
    workbook=load_workbook(path);mutator(workbook);workbook.save(path)
def mutate_json(path:Path,mutator:Callable[[dict],None])->None:
    payload=json.loads(path.read_text(encoding="utf-8"));mutator(payload);path.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
def decode_message(message)->str:
    if message.is_multipart():return "\n".join((part.get_payload(decode=True) or b"").decode(part.get_content_charset() or "utf-8",errors="replace") for part in message.walk() if part.get_content_type()=="text/plain").strip()
    return (message.get_payload(decode=True) or b"").decode(message.get_content_charset() or "utf-8",errors="replace").strip()


MBOX_HEADERS=(("from","From"),("to","To"),("cc","Cc"),("bcc","Bcc"),("subject","Subject"),("date","Date"),("message_id","Message-ID"),("status","X-EvoCowork-Status"),("order","X-EvoCowork-Order"),("source_row","X-EvoCowork-Source-Row"),("severity","X-EvoCowork-Severity"),("proposed_status","X-EvoCowork-Proposed-Status"),("osha","X-EvoCowork-OSHA-Reportable"),("workers_comp","X-EvoCowork-Workers-Comp-Required"),("email_sent","X-EvoCowork-Email-Sent"),("tracker_applied","X-EvoCowork-Tracker-Applied"),("slack_posted","X-EvoCowork-Slack-Posted"),("osha_filed","X-EvoCowork-OSHA-Filed"))
def mbox_states(path:Path)->list[dict]:
    box=mailbox.mbox(path,create=False);messages=list(box);box.close();return [{**{key:str(message.get(header,"")) for key,header in MBOX_HEADERS},"body":decode_message(message)} for message in messages]
def write_mbox_states(path:Path,states:list[dict])->None:
    chunks=[]
    for order,state in enumerate(states,1):
        message=EmailMessage(policy=SMTP)
        for key,header in MBOX_HEADERS:
            if state.get(key):message[header]=state[key]
        message.set_content(state["body"].rstrip()+"\n",charset="utf-8");buffer=BytesIO();BytesGenerator(buffer,policy=SMTP,mangle_from_=True,maxheaderlen=0).flatten(message);chunks.extend([f"From MAILER-DAEMON Wed Apr 15 09:00:{order:02d} 2026\n".encode("ascii"),buffer.getvalue().replace(b"\r\n",b"\n"),b"\n"])
    path.write_bytes(b"".join(chunks))
def mutate_message(path:Path,index:int,mutator:Callable[[dict],None])->None:
    states=mbox_states(path);mutator(states[index]);write_mbox_states(path,states)


def replace_docx_text(path:Path,old:str,new:str,*,row_token:str|None=None,min_cells:int=0)->None:
    document=Document(path);changed=False
    if row_token is None:
        for paragraph in document.paragraphs:
            if old in paragraph.text:paragraph.text=paragraph.text.replace(old,new,1);changed=True;break
    if not changed:
        for table in document.tables:
            for row in table.rows:
                if len(row.cells)<min_cells:continue
                if row_token and row_token not in " | ".join(cell.text for cell in row.cells):continue
                for cell in row.cells:
                    if old in cell.text:cell.text=cell.text.replace(old,new,1);changed=True;break
                if changed:break
            if changed:break
    if not changed:raise ValueError(f"Word找不到{old!r}")
    document.save(path)
def replace_all_docx_text(path:Path,old:str,new:str)->None:
    document=Document(path);changed=False
    for paragraph in document.paragraphs:
        if old in paragraph.text:paragraph.text=paragraph.text.replace(old,new);changed=True
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if old in cell.text:cell.text=cell.text.replace(old,new);changed=True
    if not changed:raise ValueError(f"Word找不到{old!r}")
    document.save(path)
def replace_pptx_text(path:Path,old:str,new:str,*,row_token:str|None=None,min_cells:int=0)->None:
    presentation=Presentation(path);changed=False
    for slide in presentation.slides:
        for shape in slide.shapes:
            if row_token is None and getattr(shape,"has_text_frame",False) and old in shape.text:shape.text=shape.text.replace(old,new,1);changed=True;break
            if getattr(shape,"has_table",False):
                for row in shape.table.rows:
                    if len(row.cells)<min_cells:continue
                    if row_token and row_token not in " | ".join(cell.text for cell in row.cells):continue
                    for cell in row.cells:
                        if old in cell.text:cell.text=cell.text.replace(old,new,1);changed=True;break
                    if changed:break
            if changed:break
        if changed:break
    if not changed:raise ValueError(f"PPT找不到{old!r}")
    presentation.save(path)
def drop_ppt_slide(path:Path)->None:
    presentation=Presentation(path);presentation.slides._sldIdLst.remove(presentation.slides._sldIdLst[-1]);presentation.save(path)


def build_cases()->list[Case]:
    cases:list[Case]=[]
    for name in FILES:cases.append((f"missing-{name}",lambda root,n=name:(root/n).unlink(),"R001"))
    cases.append(("extra-file",lambda root:(root/"unexpected.txt").write_text("unexpected",encoding="utf-8"),"R001"))
    corrupt={XLSX:"R001",DOCX:"R001",PPTX:"R001",MBOX:"R011",UPDATES:"R012",SLACK:"R013",EXECUTION:"R014"}
    for name in FILES:cases.append((f"corrupt-{name}",lambda root,n=name:(root/n).write_bytes(b"not-a-valid-artifact"),corrupt[name]))

    source_wrong={"A":99,"B":"Wrong Employee","C":"wrong@company.example.com","D":46000,"E":"Wrong incident","F":"Wrong Severity","G":"Wrong treatment","H":"99","I":"Wrong Status","J":"Wrong notes"}
    for row in range(2,10):
        for column,value in source_wrong.items():cases.append((f"source-{column}{row}",lambda root,c=f"{column}{row}",v=value:edit_cell(root/XLSX,"Source_Incidents",c,value=v),"R002"))
    for column,value in (("A","WRONG-ID"),("B","WRONG-THREAD"),("C","wrong@company.example.com"),("D","wrong@company.example.com"),("E","Wrong subject"),("F","Wrong date"),("G","WRONG"),("H","Wrong body"),("I","Wrong time"),("J","Wrong source")):cases.append((f"source-manual-{column}2",lambda root,c=f"{column}2",v=value:edit_cell(root/XLSX,"Source_Manual",c,value=v),"R003"))
    for row in range(2,12):
        for column,value in (("A","WRONG"),("B","WRONG"),("C","Wrong scope"),("D","Wrong requirement"),("E","wrong-source")):cases.append((f"policy-{column}{row}",lambda root,c=f"{column}{row}",v=value:edit_cell(root/XLSX,"Policy_Source",c,value=v),"R003"))
    for row in range(2,24):
        for column,value in (("A","WRONG-ID"),("B","WRONG-TYPE"),("C","Wrong entity"),("D","Wrong fact"),("E","wrong-source")):cases.append((f"evidence-{column}{row}",lambda root,c=f"{column}{row}",v=value:edit_cell(root/XLSX,"Evidence_Register",c,value=v),"R003"))

    row_criteria={2:"R005",3:"R006",4:"R007",5:"R005",6:"R007",7:"R005",8:"R006",9:"R007"}
    for row in range(2,10):
        for column in range(1,32):
            coordinate=f"{get_column_letter(column)}{row}";cases.append((f"formula-{coordinate}",lambda root,c=coordinate:edit_cell(root/XLSX,"Incident_Decisions",c,remove_formula=True),"R004"))
            criterion="R009" if column==19 else ("R008" if row in (4,6,9) and column in (14,15,16,17,20,21) else row_criteria[row])
            cases.append((f"decision-{coordinate}",lambda root,c=coordinate:edit_cell(root/XLSX,"Incident_Decisions",c,remove_formula=True,value="WRONG"),criterion))
    for row in range(2,5):
        for column in range(1,15):
            coordinate=f"{get_column_letter(column)}{row}";criterion="R009" if column==9 else "R005";cases.append((f"minor-{coordinate}",lambda root,c=coordinate:edit_cell(root/XLSX,"Minor_Routing",c,value="WRONG"),criterion))
    for row in range(2,4):
        for column in range(1,17):
            coordinate=f"{get_column_letter(column)}{row}";criterion="R009" if column==10 else "R006";cases.append((f"moderate-{coordinate}",lambda root,c=coordinate:edit_cell(root/XLSX,"Moderate_Routing",c,value="WRONG"),criterion))
    for row in range(2,5):
        for column in range(1,21):
            coordinate=f"{get_column_letter(column)}{row}";criterion="R009" if column==14 else ("R008" if column in (8,9,10,11,12,15,18) else "R007");cases.append((f"serious-{coordinate}",lambda root,c=coordinate:edit_cell(root/XLSX,"Serious_OSHA_Routing",c,value="WRONG"),criterion))
    for row,criterion in row_criteria.items():
        for column in range(1,15):
            coordinate=f"{get_column_letter(column)}{row}";cases.append((f"communication-{coordinate}",lambda root,c=coordinate:edit_cell(root/XLSX,"Communication_Register",c,value="WRONG"),criterion))
    for column in range(1,17):
        coordinate=f"{get_column_letter(column)}2";cases.append((f"slack-sheet-{coordinate}",lambda root,c=coordinate:edit_cell(root/XLSX,"Slack_Summary_Control",c,value="WRONG"),"R013"))
    for row in range(2,24):
        cases.append((f"summary-B{row}",lambda root,c=f"B{row}":edit_cell(root/XLSX,"Control_Summary",c,value=99),"R010"));cases.append((f"summary-D{row}",lambda root,c=f"D{row}":edit_cell(root/XLSX,"Control_Summary",c,value="FAIL"),"R010"))
    cases.extend([("xlsx-table-removed",lambda root:openpyxl_mutation(root/XLSX,lambda wb:wb["Source_Incidents"].tables.clear()),"R018"),("xlsx-chart-removed",lambda root:openpyxl_mutation(root/XLSX,lambda wb:setattr(wb["Control_Summary"],"_charts",[])),"R018"),("xlsx-cf-removed",lambda root:openpyxl_mutation(root/XLSX,lambda wb:setattr(wb["Incident_Decisions"],"conditional_formatting",type(wb["Incident_Decisions"].conditional_formatting)())),"R018")])

    names=["Ana Petrova","Ben Torres","Carlos Mendoza","Diana Osei","Ethan Park","Fiona Strand","George Tan","Holly Ng"]
    criteria=["R005","R006","R007","R005","R007","R005","R006","R007"]
    for index,(name,criterion) in enumerate(zip(names,criteria)):
        fields=(("from","wrong@company.example.com",criterion),("to","wrong@company.example.com",criterion),("cc","cc@company.example.com",criterion),("bcc","bcc@company.example.com",criterion),("subject","Wrong subject",criterion),("body","Wrong body",criterion),("date","Thu, 16 Apr 2026 09:00:00 +0000","R011"),("message_id","<wrong@evocowork.local>","R011"),("status","SENT","R011"),("order","99","R011"),("source_row","99","R011"),("severity","Wrong","R011"),("proposed_status","Wrong","R011"),("osha","Wrong","R011"),("workers_comp","Wrong","R011"),("email_sent","YES","R011"),("tracker_applied","YES","R011"),("slack_posted","YES","R011"),("osha_filed","YES","R011"))
        for field,value,expected in fields:cases.append((f"mbox-{index}-{field}",lambda root,i=index,f=field,v=value:mutate_message(root/MBOX,i,lambda state:state.update({f:v})),expected))
    cases.extend([("mbox-drop",lambda root:mutate_message_list(root/MBOX,lambda states:states.pop()),"R011"),("mbox-reverse",lambda root:mutate_message_list(root/MBOX,lambda states:states.reverse()),"R011")])
    for index in (2,4,7):
        cases.append((f"mbox-osha-time-{index}",lambda root,i=index:mutate_message(root/MBOX,i,lambda state:state.update({"body":state["body"].replace("exact incident time is unavailable","exact incident time is 12:00",1)})),"R008"));cases.append((f"mbox-osha-filed-{index}",lambda root,i=index:mutate_message(root/MBOX,i,lambda state:state.update({"body":state["body"].replace("OSHA filing completed: NO","OSHA filing completed: YES",1)})),"R008"))
    for index in (1,2,4,6,7):cases.append((f"mbox-workers-form-{index}",lambda root,i=index:mutate_message(root/MBOX,i,lambda state:state.update({"body":state["body"].replace("form created: NO","form created: YES",1)})),"R009"))

    global_updates=(("task_id","wrong"),("status","EXECUTED"),("frozen_time","2026-04-16T09:00:00Z"),("position_count",7),("osha_reportable_count",2),("workers_comp_required_count",4),("email_draft_count",7),("emails_sent",1),("tracker_update_draft_count",7),("tracker_updates_applied",1),("osha_reports_filed",1),("workers_comp_forms_created",1),("slack_posts",1))
    for key,value in global_updates:cases.append((f"updates-{key}",lambda root,k=key,v=value:mutate_json(root/UPDATES,lambda p:p.update({k:v})),"R012"))
    cases.extend([("updates-status-count",lambda root:mutate_json(root/UPDATES,lambda p:p["status_counts"].update({"Documented":4})),"R012"),("updates-drop",lambda root:mutate_json(root/UPDATES,lambda p:p["updates"].pop()),"R012"),("updates-reverse",lambda root:mutate_json(root/UPDATES,lambda p:p["updates"].reverse()),"R012")])
    entry_wrong={"source_row":99,"employee":"Wrong Employee","source_status":"Wrong","severity":"Wrong","proposed_validation_status":"Wrong","hospitalization_trigger":True,"amputation_trigger":True,"vision_source":"CONFIRMED_EYE_LOSS","osha_reportable":True,"legal_immediate":True,"workers_comp_required":False,"recipients":["wrong@company.example.com"],"email_subject":"Wrong","email_state":"SENT","tracker_update_applied":True,"osha_filed":True,"workers_comp_form_created":True,"slack_posted":True,"source_trace":[]}
    for index,criterion in enumerate(criteria):
        for key,value in entry_wrong.items():
            expected="R008" if key in ("hospitalization_trigger","amputation_trigger","vision_source","osha_reportable","legal_immediate","osha_filed") else ("R009" if key in ("workers_comp_required","workers_comp_form_created") else criterion)
            cases.append((f"updates-{index}-{key}",lambda root,i=index,k=key,v=value:mutate_json(root/UPDATES,lambda p:p["updates"][i].update({k:not p["updates"][i][k] if isinstance(p["updates"][i].get(k),bool) else v})),expected))

    slack_mutations=(("task_id","wrong"),("decision","POST_PUBLIC_DETAILS"),("draft_count",0),("posted_count",1),("posted",True),("state","POSTED"),("contains_employee_names",True),("contains_manager_addresses",True),("contains_medical_details",True),("contains_incident_descriptions",True),("draft_text","Ana Petrova hospitalization"),("messages_created",["posted"]),("source_trace",[]))
    for key,value in slack_mutations:cases.append((f"slack-{key}",lambda root,k=key,v=value:mutate_json(root/SLACK,lambda p:p.update({k:v})),"R013"))
    for key,value in (("incidents",7),("minor",4),("moderate",3),("serious",2),("osha_legal_immediate",2),("workers_comp_follow_up",4)):cases.append((f"slack-count-{key}",lambda root,k=key,v=value:mutate_json(root/SLACK,lambda p:p["aggregate_counts"].update({k:v})),"R013"))
    cases.extend([("slack-channel",lambda root:mutate_json(root/SLACK,lambda p:p["source_channel"].update({"id":"C_OTHER"})),"R013"),("slack-private",lambda root:mutate_json(root/SLACK,lambda p:p["source_channel"].update({"is_private":True})),"R013")])

    execution_mutations=[
        ("task-id","R014",lambda p:p.update({"task_id":"wrong"})),("status","R014",lambda p:p.update({"status":"EXECUTED"})),("retention","R014",lambda p:p.update({"retention_years":1})),("side-effect","R014",lambda p:p.update({"real_side_effect_count":1})),("email-draft","R014",lambda p:p["email"].update({"draft_count":7})),("email-sent","R014",lambda p:p["email"].update({"sent_count":1,"state":"SENT"})),("tracker-draft","R014",lambda p:p["tracker"].update({"update_draft_count":7})),("tracker-applied","R014",lambda p:p["tracker"].update({"updates_applied":1,"source_modified":True,"state":"MODIFIED"})),("slack-draft","R014",lambda p:p["slack"].update({"draft_count":0})),("slack-post","R014",lambda p:p["slack"].update({"posted_count":1,"state":"POSTED"})),("osha-count","R008",lambda p:p["osha"].update({"reportable_count":2})),("osha-legal","R008",lambda p:p["osha"].update({"legal_notification_drafts":2})),("osha-filed","R008",lambda p:p["osha"].update({"reports_filed":1})),("osha-deadline","R008",lambda p:p["osha"].update({"exact_deadlines_asserted":3})),("osha-reason","R008",lambda p:p["osha"].update({"reason":"Exact times known"})),("workers-count","R009",lambda p:p["workers_comp"].update({"follow_up_required_count":4})),("workers-form","R009",lambda p:p["workers_comp"].update({"forms_created":5})),("workers-owner","R009",lambda p:p["workers_comp"].update({"owner":"HR Ops completed"})),("trace","R014",lambda p:p.update({"source_trace":[]})),
    ]
    for name,criterion,mutator in execution_mutations:cases.append((f"execution-{name}",lambda root,m=mutator:mutate_json(root/EXECUTION,m),criterion))

    word_summary=[("Minor","3","4"),("Moderate","2","3"),("Serious","3","2")]
    for name,old,new in word_summary:cases.append((f"word-summary-{name}",lambda root,t=name,o=old,n=new:replace_docx_text(root/DOCX,o,n,row_token=t,min_cells=6),"R016"))
    for name,criterion in zip(names,criteria):cases.append((f"word-row-{name}",lambda root,t=name:replace_docx_text(root/DOCX,t,"Wrong Employee",row_token=t,min_cells=9),criterion))
    word_narratives=[("Carlos Mendoza, Ethan Park, and Holly Ng are OSHA-reportable","Only Carlos is OSHA-reportable","R008"),("Holly also has a finger amputation","Holly has no amputation","R008"),("potential loss of vision, not a confirmed loss of an eye","confirmed loss of an eye","R008"),("Exact incident time unavailable; no exact 24h cutoff asserted","Exact cutoff is 12:00","R008"),("Ben, Carlos, Ethan, George, and Holly require follow-up","Only Ben requires follow-up","R009"),("Benefits team handles claim forms separately","HR Ops submitted all forms","R009"),("0 workers' comp forms","5 workers' comp forms","R009"),("No names, manager emails, treatments, or descriptions","All names and treatments included","R013"),("Employee names, manager addresses, medical details, and incident descriptions are excluded","Employee names and medical details are included","R013")]
    for old,new,criterion in word_narratives:
        mutator=(lambda root,o=old,n=new:replace_all_docx_text(root/DOCX,o,n)) if old.startswith("Exact incident time unavailable") else (lambda root,o=old,n=new:replace_docx_text(root/DOCX,o,n))
        cases.append((f"word-narrative-{old[:14]}",mutator,criterion))
    for token,old,new,criterion in (("Incident email","0","1","R014"),("Tracker update","0","1","R014"),("Slack aggregate summary","0","1","R014"),("OSHA filing","0","1","R014"),("Workers' comp forms","0","5","R014"),("Real side effects","0","1","R014")):cases.append((f"word-control-{token}",lambda root,t=token,o=old,n=new:replace_docx_text(root/DOCX,o,n,row_token=t,min_cells=4),criterion))

    ppt_summary=[("3 Minor","4 Minor","R015"),("2 Moderate","3 Moderate","R015"),("3 Serious","2 Serious","R015"),("3 OSHA/Legal immediate","2 OSHA/Legal immediate","R015"),("5 workers' comp follow-ups","4 workers' comp follow-ups","R015"),("Exact OSHA cutoff time cannot be computed: source has dates, not incident times.","Exact OSHA cutoff is known.","R008"),("aggregate draft only; no employee names, manager addresses, treatments, or incident descriptions","public post includes employee medical details","R013")]
    for old,new,criterion in ppt_summary:cases.append((f"ppt-summary-{old[:12]}",lambda root,o=old,n=new:replace_pptx_text(root/PPTX,o,n),criterion))
    for name,criterion in zip(names,criteria):cases.append((f"ppt-row-{name}",lambda root,t=name:replace_pptx_text(root/PPTX,t,"Wrong Employee",row_token=t,min_cells=8),criterion))
    for old,new in (("Carlos: hospitalization","Carlos: no hospitalization"),("Ethan: hospitalization + potential vision loss","Ethan: no trigger"),("Holly: hospitalization + finger amputation","Holly: no trigger")):cases.append((f"ppt-osha-{old[:8]}",lambda root,o=old,n=new:replace_pptx_text(root/PPTX,o,n),"R008"))
    for old,new,criterion in (("Email 8 drafts · 0 sent","Email 8 drafts · 1 sent","R014"),("Tracker 8 drafts · 0 applied","Tracker 8 drafts · 1 applied","R014"),("Slack 1 draft · 0 posted","Slack 1 draft · 1 posted","R014"),("Names / medical details = 0","Names / medical details = 8","R014"),("OSHA reports filed = 0","OSHA reports filed = 3","R014"),("Workers' comp forms = 0","Workers' comp forms = 5","R014"),("no real side effects","real side effects executed","R014")):cases.append((f"ppt-control-{old[:10]}",lambda root,o=old,n=new:replace_pptx_text(root/PPTX,o,n),criterion))
    cases.append(("ppt-drop-slide",lambda root:drop_ppt_slide(root/PPTX),"R015"))
    return cases


def mutate_message_list(path:Path,mutator:Callable[[list[dict]],None])->None:
    states=mbox_states(path);mutator(states);write_mbox_states(path,states)


def main()->None:
    parser=argparse.ArgumentParser();parser.add_argument("--json-out",type=Path);args=parser.parse_args();judge=load_judge();cases=build_cases();results=[];failures=[]
    for name,mutator,expected_failure in cases:
        safe_name="".join(character if character.isalnum() or character in "-_" else "_" for character in name)
        with tempfile.TemporaryDirectory(prefix=f"negative-5126-{safe_name}-") as temporary:
            artifact_root=Path(temporary)/"artifacts";shutil.copytree(GOLDEN,artifact_root)
            try:
                mutator(artifact_root);result=judge.evaluate(artifact_root);detected=not result.get("criteria",{}).get(expected_failure,{}).get("passed",True);rejected=not result.get("passed",False) and float(result.get("reward",1.0))<0.8;results.append({"case":name,"expected_failure":expected_failure,"detected":detected,"rejected":rejected,"reward":result.get("reward"),"hurdle_failures":result.get("hurdle_failures",[])})
                if not detected or not rejected:failures.append(f"{name}: expected={expected_failure}, detected={detected}, reward={result.get('reward')}")
            except Exception as exc:failures.append(f"{name}: mutation/evaluation error: {type(exc).__name__}: {exc}")
    payload={"task_id":"cowork-evo/automationbench-5126-safety-incident-control-package","case_count":len(results),"declared_case_count":len(cases),"all_rejected":not failures and len(results)==len(cases),"max_reward":max((float(item["reward"]) for item in results),default=None),"failures":failures,"results":results}
    if args.json_out:args.json_out.parent.mkdir(parents=True,exist_ok=True);args.json_out.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(payload,ensure_ascii=False,indent=2))
    if failures:raise SystemExit(1)


if __name__=="__main__":main()
