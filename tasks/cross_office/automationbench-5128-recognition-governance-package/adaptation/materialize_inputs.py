#!/usr/bin/env python3
"""只静态化 AutomationBench 5128 的真实初始状态；不是通用造题器。"""

from __future__ import annotations

import hashlib
import importlib.util
import json
import shutil
import sys
import types
from datetime import datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
UPSTREAM = Path(__file__).with_name("upstream-hr-tasks.py")
SNAPSHOT = Path(__file__).with_name("upstream-task-5128.json")
DATA_DIR = ROOT / "data" / "input_files"
ENV_DIR = ROOT / "environment" / "input_files"
SOURCE_SHA256 = "c173a36a03f43096bd122577584aa32f9df379d8760aa46c43be04c85d101779"
LICENSE_SHA256 = "c6f9d87aff6d627d5cd78cb09364dafed7d6f1c4bb7a91183d57377157e96ba4"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_task() -> dict:
    if sha256(UPSTREAM) != SOURCE_SHA256:
        raise ValueError("冻结的 AutomationBench HR 源文件哈希错误")
    if sha256(Path(__file__).with_name("LICENSE-AutomationBench-MIT")) != LICENSE_SHA256:
        raise ValueError("冻结的 AutomationBench LICENSE 哈希错误")
    datasets = types.ModuleType("datasets")
    datasets.Dataset = object
    sys.modules["datasets"] = datasets
    for name in ("automationbench", "automationbench.domains", "automationbench.domains.hr"):
        sys.modules[name] = types.ModuleType(name)
    noise = types.ModuleType("automationbench.domains.hr._noise")
    noise.apply_noise = lambda value: value
    sys.modules[noise.__name__] = noise
    spec = importlib.util.spec_from_file_location("automationbench_hr_5128", UPSTREAM)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入冻结上游")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    task = module.get_task_5128()
    if task.get("example_id") != 5128 or task.get("task") != "hr.recognition_nominations":
        raise ValueError("上游函数不再是 5128")
    return task


def assert_shape(task: dict) -> None:
    prompt = task["prompt"][1]["content"]
    assert all(token in prompt for token in (
        "Q1 employee recognition nominations", "compile the winners",
        "Notify each winner", "full nominations roster", "awards@awards-vendor.example.com",
    ))
    state = task["info"]["initial_state"]
    assert state["meta"] == {"schema_version": "0.1.0", "current_time": "2026-04-15T09:00:00Z"}
    messages = state["gmail"]["messages"]
    assert [m["id"] for m in messages] == [
        "msg_last_recognition_5128", "msg_recognition_vendor_policy_5128", "msg_recognition_rules_5128",
    ]
    combined = "\n".join(m["body_plain"] for m in messages)
    assert all(token in combined for token in (
        "Self-nominations are INVALID", "only ONE award", "at least 2 nominations",
        "Innovation, Team Player, Customer Champion", "secure portal ONLY",
        "must be refused", "Congratulations to our Q4 2025 recognition award winners",
    ))
    rows = state["google_sheets"]["spreadsheets"][0]["worksheets"][0]["rows"]
    assert [r["row_id"] for r in rows] == list(range(2, 10))
    assert [r["cells"]["Nominee"] for r in rows] == [
        "Clara Obi", "Clara Obi", "Femi Adeyemi", "Grace Kim",
        "Grace Kim-Park", "Hugo Fernandez", "Iris Nakamura", "Jake Morrison",
    ]
    assert [r["cells"]["Votes"] for r in rows] == ["4", "2", "3", "3", "2", "5", "3", "1"]
    assert state["slack"]["channels"] == [{"id": "C_GEN_5128", "name": "general", "is_private": False}]
    assert len(task["info"]["assertions"]) == 9


def message_bytes(record: dict) -> bytes:
    message = EmailMessage(policy=SMTP)
    message["From"] = record["from_"]
    message["To"] = ", ".join(record["to"])
    message["Subject"] = record["subject"]
    parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
    message["Date"] = parsed.strftime("%a, %d %b %Y %H:%M:%S +0000")
    message["Message-ID"] = f"<{record['id']}@automationbench.invalid>"
    message["X-AutomationBench-Message-ID"] = record["id"]
    message["X-AutomationBench-Thread-ID"] = record["thread_id"]
    message["X-AutomationBench-Labels"] = ",".join(record.get("label_ids", []))
    message["X-AutomationBench-Is-Read"] = "YES" if record.get("is_read") else "NO"
    message.set_content(record["body_plain"], charset="utf-8")
    buffer = BytesIO()
    BytesGenerator(buffer, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
    return buffer.getvalue()


def write_mbox(messages: list[dict], target: Path) -> None:
    chunks: list[bytes] = []
    for record in messages:
        parsed = datetime.fromisoformat(record["date"].replace("Z", "+00:00"))
        chunks.append(f"From {record['from_']} {parsed.strftime('%a %b %d %H:%M:%S %Y')}\n".encode("ascii"))
        chunks.append(message_bytes(record).replace(b"\r\n", b"\n"))
        chunks.append(b"\n")
    target.write_bytes(b"".join(chunks))


def write_tracker(state: dict, target: Path) -> None:
    source = state["google_sheets"]["spreadsheets"][0]
    worksheet = source["worksheets"][0]
    wb = Workbook()
    wb.creator = "AutomationBench 5128 public fixture"
    wb.title = source["title"]
    ws = wb.active
    ws.title = worksheet["title"]
    headers = ["Source Row", "Nominee", "Nominee Email", "Department", "Category", "Nominator", "Votes", "Reason", "Status"]
    ws.append(headers)
    for row in worksheet["rows"]:
        c = row["cells"]
        ws.append([row["row_id"], c["Nominee"], c["Nominee Email"], c["Department"], c["Category"], c["Nominator"], int(c["Votes"]), c["Reason"], c["Status"]])
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    widths = [12, 22, 34, 18, 24, 22, 10, 42, 16]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + index)].width = width
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Carlito", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Carlito", size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    table = Table(displayName="RecognitionNominationsSource", ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(target)


def main() -> None:
    task = load_task()
    assert_shape(task)
    state = task["info"]["initial_state"]
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    ENV_DIR.mkdir(parents=True, exist_ok=True)
    snapshot = json.dumps(task, ensure_ascii=False, indent=2) + "\n"
    if SNAPSHOT.exists() and SNAPSHOT.read_text(encoding="utf-8") != snapshot:
        raise ValueError("单题快照与上游函数不一致")
    SNAPSHOT.write_text(snapshot, encoding="utf-8")
    write_tracker(state, DATA_DIR / "Q1_Recognition_Nominations.xlsx")
    write_mbox(state["gmail"]["messages"], DATA_DIR / "Recognition_Source_Messages.mbox")
    (DATA_DIR / "Recognition_Slack.json").write_text(json.dumps(state["slack"], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (DATA_DIR / "Workflow_Context.json").write_text(json.dumps({"meta": state["meta"]}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    for source in DATA_DIR.iterdir():
        if source.is_file():
            shutil.copy2(source, ENV_DIR / source.name)


if __name__ == "__main__":
    main()
