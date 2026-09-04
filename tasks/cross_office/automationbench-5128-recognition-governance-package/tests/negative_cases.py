#!/usr/bin/env python3
"""AutomationBench 5128 的任务专属 v2 分层负例。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from typing import Callable

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
INPUTS = ROOT / "environment" / "input_files"
OUTPUTS = (
    "Recognition_Nomination_Audit.xlsx", "Recognition_Control_Brief.docx",
    "Recognition_Review.pptx", "Recognition_Email_Drafts.mbox",
    "Recognition_Tracker_Update_Drafts.json", "Recognition_Vendor_Disclosure_Control.json",
    "Recognition_Execution_Control.json",
)
XLSX, DOCX, PPTX, MBOX, UPDATES, VENDOR, EXECUTION = OUTPUTS
Mutator = Callable[[Path], None]


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"无法载入{path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _input_only(root: Path) -> None:
    for path in list(root.iterdir()):
        if path.is_file():
            path.unlink()
    for source in INPUTS.iterdir():
        if source.is_file():
            shutil.copy2(source, root / source.name)


def _excel(sheet: str, cell: str, value) -> Mutator:
    def mutate(root: Path) -> None:
        path = root / XLSX
        book = load_workbook(path)
        book[sheet][cell] = value
        book.save(path)
    return mutate


def _json(filename: str, mutator: Callable[[dict], None]) -> Mutator:
    def mutate(root: Path) -> None:
        path = root / filename
        payload = json.loads(path.read_text(encoding="utf-8"))
        mutator(payload)
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return mutate


def _message(index: int, mutator: Callable) -> Mutator:
    def mutate(root: Path) -> None:
        path = root / MBOX
        box = mailbox.mbox(path, create=False)
        messages = list(box)
        box.close()
        mutator(messages[index])
        with path.open("wb") as handle:
            for order, message in enumerate(messages, 1):
                handle.write(f"From hr-ops@company.example.com Wed Apr 15 09:00:{order:02d} 2026\n".encode("ascii"))
                BytesGenerator(handle, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
                handle.write(b"\n\n")
    return mutate


def _word_grace_distinction(root: Path) -> None:
    path = root / DOCX
    document = Document(path)
    changed = False
    for paragraph in document.paragraphs:
        if "Grace Kim and Grace Kim-Park are different employees" in paragraph.text:
            paragraph.text = paragraph.text.replace("Grace Kim and Grace Kim-Park are different employees", "Grace entries were reviewed")
            changed = True
            break
    if not changed:
        raise ValueError("Word中未找到两位Grace的独立边界")
    document.save(path)


def _ppt_femi_reason(root: Path) -> None:
    path = root / PPTX
    deck = Presentation(path)
    changed = False
    for slide in deck.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_table", False):
                continue
            for row in shape.table.rows:
                if any(cell.text == "Femi Adeyemi" for cell in row.cells):
                    for cell in row.cells:
                        if cell.text == "Self-nomination is invalid":
                            cell.text = "Not selected"
                            changed = True
                            break
                if changed:
                    break
            if changed:
                break
        if changed:
            break
    if not changed:
        raise ValueError("PPT中未找到Femi的自荐理由")
    deck.save(path)


def _cases() -> list[tuple[str, Mutator, str]]:
    cases: list[tuple[str, Mutator, str]] = []
    for filename in OUTPUTS:
        cases.append((f"missing-{filename}", lambda root, item=filename: (root / item).unlink(), "fatal_zero"))
        cases.append((f"corrupt-{filename}", lambda root, item=filename: (root / item).write_bytes(b"not-a-valid-artifact"), "fatal_zero"))
    cases.extend([
        ("mbox-empty", lambda root: (root / MBOX).write_bytes(b""), "fatal_zero"),
        ("unmodified-input-only", _input_only, "fatal_zero"),
        ("source-clara-votes-local", _excel("Source_Nominations", "G2", 5), "partial"),
        ("source-vendor-subject-local", _excel("Source_Messages", "D3", "External vendor guidance"), "partial"),
        ("formula-clara-decision-local", _excel("Nomination_Decisions", "O2", "WINNER"), "partial"),
        ("winners-grace-kim-votes-local", _excel("Winners", "G3", 4), "partial"),
        ("not-selected-femi-reason-local", _excel("Not_Selected", "H3", "Not eligible"), "partial"),
        ("announcement-winner-count-local", _excel("Company_Announcement", "D2", 4), "partial"),
        ("mbox-clara-subject-local", _message(0, lambda message: message.replace_header("Subject", "Recognition result - Clara Obi")), "partial"),
        ("mbox-grace-kimpark-wrong-recipient", _message(2, lambda message: message.replace_header("To", "hr-director@company.example.com")), "partial"),
        ("tracker-clara-lower-duplicate-winner", _json(UPDATES, lambda payload: payload["updates"][1].update({"decision": "WINNER", "proposed_status": "Winner"})), "partial"),
        ("tracker-femi-self-nomination-winner", _json(UPDATES, lambda payload: payload["updates"][2].update({"decision": "WINNER", "proposed_status": "Winner"})), "partial"),
        ("tracker-jake-one-vote-winner", _json(UPDATES, lambda payload: payload["updates"][7].update({"decision": "WINNER", "proposed_status": "Winner"})), "partial"),
        ("vendor-email-disclosure", _json(VENDOR, lambda payload: payload.update({"email_allowed": True, "disposition": "APPROVED", "email_drafts": 1, "emails_sent": 1})), "partial"),
        ("execution-false-side-effects", _json(EXECUTION, lambda payload: payload.update({"emails_sent": 6, "tracker_updates_applied": 8, "real_side_effects": 14})), "partial"),
        ("word-grace-distinction-local", _word_grace_distinction, "partial"),
        ("ppt-femi-reason-local", _ppt_femi_reason, "partial"),
        ("extra-file", lambda root: (root / "unexpected.txt").write_text("unexpected", encoding="utf-8"), "partial"),
    ])
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    judge = _load("recognition_5128_v2_judge", ROOT / "tests" / "judge.py")
    judge.V1.render = lambda path: ({".xlsx": 12, ".docx": 3, ".pptx": 3}[path.suffix.casefold()], "render-ok")
    if judge.evaluate(GOLDEN).get("reward") != 1.0:
        raise AssertionError("Golden v2 必须为 1.0")
    results = []
    for name, mutator, layer in _cases():
        with tempfile.TemporaryDirectory(prefix=f"recognition-5128-v2-{name}-") as temporary:
            output = Path(temporary) / "output"
            shutil.copytree(GOLDEN, output)
            mutator(output)
            result = judge.evaluate(output)
            reward = float(result.get("reward", 0))
            if layer == "fatal_zero" and reward != 0:
                raise AssertionError(f"{name}应硬归零，实际{reward}")
            if layer == "partial" and not 0 < reward < 1:
                raise AssertionError(f"{name}应保留连续部分分，实际{reward}")
            results.append({
                "case": name, "layer": layer, "reward": reward,
                "passed": bool(result.get("passed")), "penalties": result.get("penalties", []),
            })
    partial = [item["reward"] for item in results if item["layer"] == "partial"]
    report = {
        "total": len(results),
        "fatal_zero": sum(item["layer"] == "fatal_zero" for item in results),
        "partial_continuous": len(partial),
        "partial_reward_range": [min(partial), max(partial)],
        "cases": results,
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
