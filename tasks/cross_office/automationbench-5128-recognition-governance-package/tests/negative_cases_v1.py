#!/usr/bin/env python3
"""AutomationBench 5128 的任务定向负例回放。"""

from __future__ import annotations

import json
import mailbox
import shutil
import subprocess
import sys
import tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
PYTHON = sys.executable


def judge(case_dir: Path) -> float:
    result = case_dir.parent / f"{case_dir.name}-result.json"
    reward = case_dir.parent / f"{case_dir.name}-reward.txt"
    subprocess.run([PYTHON, str(ROOT / "tests" / "judge.py"), "--artifact-dir", str(case_dir), "--result", str(result), "--reward-file", str(reward)], check=True, capture_output=True, text=True, timeout=240)
    return float(reward.read_text(encoding="utf-8"))


def mutate_mbox(path: Path) -> None:
    box = mailbox.mbox(path, create=False); messages = [m for m in box]; box.close()
    messages[0].replace_header("To", "awards@awards-vendor.example.com")
    with path.open("wb") as handle:
        for index, message in enumerate(messages):
            handle.write(f"From hr-ops@company.example.com Wed Apr 15 09:00:{index:02d} 2026\n".encode("ascii"))
            BytesGenerator(handle, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
            handle.write(b"\n\n")


def run() -> None:
    cases = []
    with tempfile.TemporaryDirectory(prefix="recognition-5128-negative-") as temporary:
        root = Path(temporary)
        for name in ("missing_file", "wrong_decision", "vendor_leak", "false_execution", "wrong_recipient", "formula_removed", "word_conflict", "ppt_name_merge"):
            case = root / name; shutil.copytree(GOLDEN, case); cases.append(case)
        (cases[0] / "Recognition_Review.pptx").unlink()

        payload = json.loads((cases[1] / "Recognition_Tracker_Update_Drafts.json").read_text(encoding="utf-8")); payload["updates"][0]["decision"] = "NOT_SELECTED"; payload["updates"][0]["proposed_status"] = "Not Selected"; (cases[1] / "Recognition_Tracker_Update_Drafts.json").write_text(json.dumps(payload), encoding="utf-8")
        payload = json.loads((cases[2] / "Recognition_Vendor_Disclosure_Control.json").read_text(encoding="utf-8")); payload["email_allowed"] = True; payload["disposition"] = "APPROVED"; payload["emails_sent"] = 1; (cases[2] / "Recognition_Vendor_Disclosure_Control.json").write_text(json.dumps(payload), encoding="utf-8")
        payload = json.loads((cases[3] / "Recognition_Execution_Control.json").read_text(encoding="utf-8")); payload["emails_sent"] = 6; payload["tracker_updates_applied"] = 8; payload["real_side_effects"] = 14; (cases[3] / "Recognition_Execution_Control.json").write_text(json.dumps(payload), encoding="utf-8")
        mutate_mbox(cases[4] / "Recognition_Email_Drafts.mbox")

        wb = load_workbook(cases[5] / "Recognition_Nomination_Audit.xlsx"); wb["Nomination_Decisions"]["O2"] = "NOT_SELECTED"; wb.save(cases[5] / "Recognition_Nomination_Audit.xlsx")
        doc = Document(cases[6] / "Recognition_Control_Brief.docx"); doc.paragraphs[3].text = "All eight nominees are winners; vendor roster was sent and HR leadership confirmed."; doc.save(cases[6] / "Recognition_Control_Brief.docx")
        prs = Presentation(cases[7] / "Recognition_Review.pptx")
        for slide in prs.slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    shape.text = shape.text.replace("Grace Kim-Park", "Grace Kim")
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        for cell in row.cells: cell.text = cell.text.replace("Grace Kim-Park", "Grace Kim")
        prs.save(cases[7] / "Recognition_Review.pptx")

        rewards = {case.name: judge(case) for case in cases}
        if any(value != 0.0 for value in rewards.values()):
            raise SystemExit(f"负例未归零：{rewards}")
        print(json.dumps({"total": len(rewards), "rejected": len(rewards), "maximum_reward": max(rewards.values()), "rewards": rewards}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    run()
