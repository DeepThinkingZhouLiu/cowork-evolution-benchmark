#!/usr/bin/env python3
"""AutomationBench 5129 的 v2 连续评分入口；原严格口径保存在 judge_v1.py。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import mailbox
import re
import sys
from email.header import decode_header, make_header
from email.utils import getaddresses
from pathlib import Path
from typing import Any, Callable

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


TASK_ID = "cowork-evo/automationbench-5129-rto-compliance-control-package"
RUBRIC_VERSION = "automationbench-5129-rto-v0.2"
JUDGE_VERSION = "automationbench-5129-criterion-partial-v2.0"
XLSX = "RTO_Compliance_Audit.xlsx"
DOCX = "RTO_Control_Brief.docx"
PPTX = "RTO_Review.pptx"
EMPLOYEE = "RTO_Employee_Reminder_Drafts.mbox"
MANAGER = "RTO_Manager_Notification_Drafts.mbox"
SLACK = "RTO_Slack_Summary_Draft.json"
EXECUTION = "RTO_Execution_Control.json"
REQUIRED = (XLSX, DOCX, PPTX, EMPLOYEE, MANAGER, SLACK, EXECUTION)
POSITIVE_WEIGHTS = {
    "R001": 5, "R002": 8, "R003": 7, "R004": 12, "R005": 12,
    "R006": 9, "R007": 9, "R008": 7, "R009": 7, "R010": 7,
    "R011": 7, "R012": 4, "R013": 6,
}
PENALTY_WEIGHTS = {"P001": -35, "P002": -25, "P003": -35, "P004": -10}


def _load_v1():
    path = Path(__file__).with_name("judge_v1.py")
    if str(path.parent) not in sys.path:
        sys.path.insert(0, str(path.parent))
    spec = importlib.util.spec_from_file_location("automationbench_5129_judge_v1", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v1 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


V1 = _load_v1()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _rows(sheet, start: int, end: int, columns: int) -> list[tuple[Any, ...]]:
    return [tuple(sheet.cell(row, column).value for column in range(1, columns + 1)) for row in range(start, end + 1)]


def _messages(path: Path) -> list:
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    return messages


def _header(message, name: str) -> str:
    value = str(message.get(name, ""))
    try:
        return str(make_header(decode_header(value)))
    except Exception:
        return value


def _addresses(message, name: str) -> list[str]:
    return [address.casefold() for _, address in getaddresses(message.get_all(name, [])) if address]


def _criterion(key: str, score: float, evidence: str, earned: float = 0, total: float = 1) -> dict[str, Any]:
    return {
        "criterion_id": key,
        "score": round(max(0.0, min(1.0, score)), 6),
        "evidence": evidence,
        "method": "deterministic",
        "raw": {"earned_units": round(earned, 6), "total_units": round(total, 6)},
    }


def _fatal(errors: list[str], hashes: list[str]) -> dict[str, Any]:
    keys = (*POSITIVE_WEIGHTS, *PENALTY_WEIGHTS)
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes or ["0" * 64],
        "criterion_results": [_criterion(key, 0, "；".join(errors), 0, 0) for key in keys],
        "reward": 0.0,
        "passed": False,
        "hurdle_failures": ["R001"],
        "penalties": [],
    }


def _core(artifact: Path) -> tuple[bool, list[str]]:
    errors: list[str] = []
    if not artifact.is_dir():
        return False, ["输出目录不存在"]
    for name in REQUIRED:
        path = artifact / name
        if not path.is_file() or path.stat().st_size == 0:
            errors.append(f"核心交付缺失或为空：{name}")
            continue
        try:
            if name == XLSX:
                book = load_workbook(path, read_only=True, data_only=False)
                book.close()
            elif name == DOCX:
                Document(path)
            elif name == PPTX:
                if not Presentation(path).slides:
                    raise ValueError("PPT 无页面")
            elif name in (EMPLOYEE, MANAGER):
                messages = _messages(path)
                if not messages or not all(message.get("To") and message.get("Subject") for message in messages):
                    raise ValueError("MBOX 无有效草稿")
            elif not isinstance(json.loads(path.read_text(encoding="utf-8")), dict):
                raise ValueError("JSON 顶层不是对象")
        except Exception as exc:
            errors.append(f"核心交付无法解析：{name}: {exc}")
    return not errors, errors


def _ratio(findings: list[tuple[bool, str]]) -> tuple[float, list[str], int, int]:
    earned = sum(bool(ok) for ok, _ in findings)
    return earned / len(findings), [label for ok, label in findings if not ok], earned, len(findings)


def _table_findings(seen: list[tuple[Any, ...]], wanted: list[tuple[Any, ...]], label: str) -> list[tuple[bool, str]]:
    findings = [(len(seen) == len(wanted), f"{label}行数错误")]
    for row_index, expected_row in enumerate(wanted):
        observed = seen[row_index] if row_index < len(seen) else ()
        for column, value in enumerate(expected_row):
            findings.append((column < len(observed) and observed[column] == value, f"{label}第{row_index + 2}行字段{column + 1}错误"))
    return findings


def _leaf_findings(observed: Any, expected: Any, path: str) -> list[tuple[bool, str]]:
    if isinstance(expected, dict):
        source = observed if isinstance(observed, dict) else {}
        findings = [(isinstance(observed, dict), f"{path}应为对象")]
        for key, value in expected.items():
            findings.extend(_leaf_findings(source.get(key), value, f"{path}.{key}"))
        return findings
    if isinstance(expected, list):
        source = observed if isinstance(observed, list) else []
        findings = [(isinstance(observed, list) and len(observed) == len(expected), f"{path}列表长度错误")]
        for index, value in enumerate(expected):
            findings.extend(_leaf_findings(source[index] if index < len(source) else None, value, f"{path}[{index}]"))
        return findings
    return [(observed == expected, f"{path}应为{expected!r}")]


def _source_messages() -> list[tuple[Any, ...]]:
    messages = _messages(V1.inputs() / "RTO_Policy_Messages.mbox")
    return [(str(m.get("X-AutomationBench-Message-ID")), str(m.get("From")), _header(m, "Subject"), V1.body_text(m)) for m in messages]


def _source_positive(artifact: Path):
    book = load_workbook(artifact / XLSX, data_only=True)
    seen = _rows(book["Source_Attendance"], 2, 9, 8)
    book.close()
    findings = _table_findings(seen, list(V1.SOURCE), "Source_Attendance")
    for filename, digest in V1.INPUT_HASHES.items():
        path = V1.inputs() / filename
        findings.append((path.is_file() and _sha256(path) == digest, f"冻结输入{filename}哈希错误"))
    return _ratio(findings)


def _trace_positive(artifact: Path):
    policy_rows = [
        ("RTO-01", "Non-remote employees require 3 in-office days", "COO"),
        ("RTO-02", "Engineering requires 2 days under approved override", "VP Engineering + COO approval"),
        ("RTO-03", "Fully remote employees are exempt", "COO"),
        ("RTO-04", "Approved WFH exception in source Notes prevents flagging", "Source attendance exception evidence"),
        ("RTO-05", "External competitor's 1-day suggestion is non-authoritative", "External distractor"),
    ]
    evidence_rows = [(name, "Frozen 5129 source", digest) for name, digest in V1.INPUT_HASHES.items()]
    book = load_workbook(artifact / XLSX, data_only=True)
    findings: list[tuple[bool, str]] = []
    for sheet_name, wanted, columns in (
        ("Source_Messages", _source_messages(), 4),
        ("Policy_Rules", policy_rows, 3),
        ("Evidence_Register", evidence_rows, 3),
    ):
        seen = _rows(book[sheet_name], 2, 1 + len(wanted), columns)
        findings.extend(_table_findings(seen, wanted, sheet_name))
    book.close()
    return _ratio(findings)


def _formula_tokens(row: int, column: int) -> tuple[str, ...]:
    if column <= 8:
        return (f"Source_Attendance!{'ABCDEFGH'[column - 1]}{row}",)
    return {
        9: (f"F{row}", "Remote", f"E{row}", "Engineering", "2", "3"),
        10: (f"F{row}", "Remote", "YES", "NO"),
        11: ("SEARCH", "Approved WFH exception", f"H{row}", "YES", "NO"),
        12: ("MAX", f"I{row}", f"G{row}", "0"),
        13: (f"J{row}", "EXEMPT_REMOTE", f"K{row}", "EXEMPT_APPROVED_EXCEPTION", f"G{row}", f"I{row}", "COMPLIANT", "NON_COMPLIANT"),
        14: (f"M{row}", "Fully remote employees are exempt", "Approved WFH exception in source Notes", f"G{row}", f"I{row}"),
        15: (f"M{row}", "NON_COMPLIANT", "YES", "NO"),
        16: (f"O{row}", f"C{row}"),
        17: (f"O{row}", f"D{row}"),
        18: (f"O{row}", "DRAFT_NOT_SENT", "NOT_REQUIRED"),
        19: (f"O{row}", "DRAFT_NOT_SENT", "NOT_REQUIRED"),
        20: ('"NO"',),
        21: ("GENERAL_POLICY", "ENGINEERING_OVERRIDE", "REMOTE_EXEMPTION", "SOURCE_NOTE_EXCEPTION"),
        22: ("source_row=", f"A{row}", "employee=", f"B{row}", "days=", f"G{row}", "required=", f"I{row}", "outcome=", f"M{row}"),
    }[column]


def _formula_positive(artifact: Path):
    book = load_workbook(artifact / XLSX, data_only=False)
    sheet = book["Compliance_Decisions"]
    formulas = [cell for row in sheet.iter_rows(min_row=2, max_row=9, min_col=1, max_col=22) for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")]
    findings: list[tuple[bool, str]] = [(len(formulas) == 176, f"Compliance_Decisions公式数应为176，实际{len(formulas)}")]
    for row in range(2, 10):
        for column in range(1, 23):
            cell = sheet.cell(row, column)
            value = str(cell.value or "")
            tokens = _formula_tokens(row, column)
            findings.append((value.startswith("=") and all(token.casefold() in value.casefold() for token in tokens), f"Compliance_Decisions!{cell.coordinate}公式依赖错误"))
    book.close()
    return _ratio(findings)


def _decision_positive(artifact: Path):
    targets = {
        "Compliant": ([(2, "Alice Park", "Engineering", "Hybrid", 2, 2, 0, "COMPLIANT", "2 in-office days meets requirement 2"), (4, "Carol Diaz", "Marketing", "Hybrid", 3, 3, 0, "COMPLIANT", "3 in-office days meets requirement 3"), (8, "Grace Okonkwo", "Sales", "Hybrid", 4, 3, 0, "COMPLIANT", "4 in-office days meets requirement 3")], 9),
        "Non_Compliant": ([(3, "Bob Chen", "Sales", "Hybrid", 2, 3, 1, "NON_COMPLIANT", "2 in-office days is below requirement 3"), (5, "Dave Kim", "Engineering", "Hybrid", 1, 2, 1, "NON_COMPLIANT", "1 in-office days is below requirement 2"), (9, "Hugo Fernandez", "Sales", "Hybrid", 2, 3, 1, "NON_COMPLIANT", "2 in-office days is below requirement 3")], 9),
        "Exempt_Or_Exception": ([(6, "Elena Vasquez", "EXEMPT_REMOTE", "Fully remote employees are exempt", 0, 0), (7, "Femi Adeyemi", "EXEMPT_APPROVED_EXCEPTION", "Approved WFH exception in source Notes", 1, 3)], 6),
        "Employee_Communications": ([(1, "Bob Chen", "bob.chen@company.example.com", 2, 3, 1, "DRAFT_NOT_SENT"), (2, "Dave Kim", "dave.kim@company.example.com", 1, 2, 1, "DRAFT_NOT_SENT"), (3, "Hugo Fernandez", "hugo.fernandez@company.example.com", 2, 3, 1, "DRAFT_NOT_SENT")], 7),
        "Manager_Communications": ([(1, "mgr-sales@company.example.com", "Bob Chen", 2, 3, 1, "DRAFT_NOT_SENT"), (2, "mgr-eng@company.example.com", "Dave Kim", 1, 2, 1, "DRAFT_NOT_SENT"), (3, "mgr-sales@company.example.com", "Hugo Fernandez", 2, 3, 1, "DRAFT_NOT_SENT")], 7),
        "Control_Summary": ([("Source employees", 8, "Complete"), ("Compliant", 3, "Alice; Carol; Grace"), ("Non-compliant", 3, "Bob; Dave; Hugo"), ("Remote exempt", 1, "Elena"), ("Approved exception", 1, "Femi"), ("Employee reminders", 3, "Draft only"), ("Manager notifications", 3, "Draft only"), ("Slack summaries", 1, "Draft only"), ("Real side effects", 0, "None")], 3),
    }
    expected_sheets = ["Source_Attendance", "Source_Messages", "Policy_Rules", "Evidence_Register", "Compliance_Decisions", "Compliant", "Non_Compliant", "Exempt_Or_Exception", "Employee_Communications", "Manager_Communications", "Control_Summary"]
    book = load_workbook(artifact / XLSX, data_only=True)
    findings: list[tuple[bool, str]] = [(book.sheetnames == expected_sheets, "Excel十一张工作表名称或顺序错误")]
    for sheet_name, (wanted, columns) in targets.items():
        seen = _rows(book[sheet_name], 2, 1 + len(wanted), columns)
        findings.extend(_table_findings(seen, wanted, sheet_name))
    book.close()
    return _ratio(findings)


def _mail_positive(artifact: Path, manager: bool):
    filename = MANAGER if manager else EMPLOYEE
    messages = _messages(artifact / filename)
    people = [
        ("Bob Chen", "bob.chen@company.example.com", "mgr-sales@company.example.com", 2, 3, 1),
        ("Dave Kim", "dave.kim@company.example.com", "mgr-eng@company.example.com", 1, 2, 1),
        ("Hugo Fernandez", "hugo.fernandez@company.example.com", "mgr-sales@company.example.com", 2, 3, 1),
    ]
    findings: list[tuple[bool, str]] = [(len(messages) == 3, f"{'经理' if manager else '员工'}MBOX不是三封逐人草稿")]
    for order, (person, employee_address, manager_address, days, required, gap) in enumerate(people, 1):
        message = messages[order - 1] if order <= len(messages) else None
        if message is None:
            findings.extend([(False, f"{person}缺少{'经理' if manager else '员工'}草稿字段") for _ in range(16)])
            continue
        recipient = manager_address if manager else employee_address
        body = V1.body_text(message)
        findings.extend([
            (_header(message, "From") == "hr-ops@company.example.com", f"{person}发件人错误"),
            (_addresses(message, "To") == [recipient], f"{person}收件人错误"),
            ((person in _header(message, "Subject")) if manager else (_header(message, "Subject") == "Weekly RTO compliance reminder"), f"{person}主题错误"),
            (_header(message, "X-EvoCowork-State") == "DRAFT_NOT_SENT", f"{person}草稿状态错误"),
            (_header(message, "X-EvoCowork-Type") == ("MANAGER_NOTIFICATION" if manager else "EMPLOYEE_REMINDER"), f"{person}邮件类型错误"),
            (not _addresses(message, "Cc") and not _addresses(message, "Bcc"), f"{person}不应有Cc/Bcc"),
            (not message.is_multipart(), f"{person}不应有附件"),
        ])
        for token in (person, "April 7–11", str(days), str(required), str(gap), "NON_COMPLIANT", "DRAFT_NOT_SENT", "no message has been sent"):
            findings.append((token in body, f"{person}正文缺少{token}"))
    return _ratio(findings)


def _expected_slack() -> dict[str, Any]:
    return {"channel":"hr-ops", "channel_id":"C_HROPS_5129", "is_private":False, "state":"DRAFT_NOT_POSTED", "draft_count":1, "posted_count":0, "text":"Weekly RTO compliance review: 8 employees reviewed; 3 non-compliant; 3 compliant; 1 fully remote exempt; 1 approved WFH exception. Follow-up drafts prepared; nothing sent or posted.", "non_compliant_count":3, "names_included":0, "real_side_effects":0}


def _expected_execution() -> dict[str, Any]:
    return {"task_id":"automationbench-5129", "employees_reviewed":8, "compliant":3, "non_compliant":3, "remote_exempt":1, "approved_exception":1, "employee_reminder_drafts":3, "manager_notification_drafts":3, "email_drafts_total":6, "emails_sent":0, "slack_drafts":1, "slack_posts":0, "tracker_update_drafts":0, "tracker_updates_applied":0, "real_side_effects":0}


def _json_positive(artifact: Path, filename: str, expected: dict[str, Any], label: str):
    return _ratio(_leaf_findings(json.loads((artifact / filename).read_text(encoding="utf-8")), expected, label))


def _slack_positive(artifact: Path):
    return _json_positive(artifact, SLACK, _expected_slack(), "RTO_Slack_Summary_Draft")


def _execution_positive(artifact: Path):
    return _json_positive(artifact, EXECUTION, _expected_execution(), "RTO_Execution_Control")


def _word_positive(artifact: Path):
    document = Document(artifact / DOCX)
    rows = [tuple(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
    text = "\n".join(paragraph.text for paragraph in document.paragraphs) + "\n" + "\n".join(" | ".join(row) for row in rows)
    findings: list[tuple[bool, str]] = [(len(document.sections) == 3, "Word不是三分节"), (len(document.tables) == 3, "Word不是三张原生表")]
    for row in [("Compliant", "3", "Alice Park; Carol Diaz; Grace Okonkwo"), ("Non-compliant", "3", "Bob Chen; Dave Kim; Hugo Fernandez"), ("Remote exempt", "1", "Elena Vasquez"), ("Approved exception", "1", "Femi Adeyemi")]:
        findings.append((row in rows, f"Word总览缺少{row[0]}"))
    for source in V1.SOURCE:
        required, _, outcome = V1.EXPECTED[source[0]]
        reason = "Fully remote employees are exempt" if source[0] == 6 else ("Approved WFH exception in source Notes" if source[0] == 7 else (f"{source[6]} in-office days meets requirement {required}" if outcome == "COMPLIANT" else f"{source[6]} in-office days is below requirement {required}"))
        row = (str(source[0]), source[1], source[4], source[5], str(source[6]), str(required), outcome, reason)
        findings.append((row in rows, f"Word缺少{source[1]}行级结论"))
    for row in [("Employee reminder", "3", "0", "Bob; Dave; Hugo"), ("Manager notification", "3", "0", "One per non-compliant employee"), ("Slack #hr-ops summary", "1", "0", "Aggregate count 3"), ("Tracker update", "0", "0", "Not requested"), ("Real side effects", "0", "0", "Offline only")]:
        findings.append((row in rows, f"Word执行控制缺少{row[0]}"))
    for token in ("Engineering two-day override", "competitor's one-day suggestion is non-authoritative", "All communications are DRAFT_NOT_SENT or DRAFT_NOT_POSTED", "no policy was changed"):
        findings.append((token.casefold() in text.casefold(), f"Word缺少{token}"))
    return _ratio(findings)


def _ppt_content(artifact: Path):
    deck = Presentation(artifact / PPTX)
    texts: list[str] = []
    rows_by_slide: list[list[tuple[str, ...]]] = []
    tables: list[int] = []
    pictures = 0
    for slide in deck.slides:
        parts: list[str] = []
        rows: list[tuple[str, ...]] = []
        table_count = 0
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                parts.append(shape.text.strip())
            if getattr(shape, "has_table", False):
                table_count += 1
                rows.extend(tuple(cell.text for cell in row.cells) for row in shape.table.rows)
        texts.append("\n".join(parts))
        rows_by_slide.append(rows)
        tables.append(table_count)
    return texts, rows_by_slide, tables, pictures


def _ppt_positive(artifact: Path):
    texts, rows_by_slide, tables, pictures = _ppt_content(artifact)
    findings: list[tuple[bool, str]] = [(len(texts) == 3, "PPT不是三页"), (tables == [1, 1, 1], "PPT不是每页一张原生表"), (pictures == 0, "PPT使用图片替代关键内容")]
    for row in [("COMPLIANT", "3", "Alice; Carol; Grace"), ("NON_COMPLIANT", "3", "Bob; Dave; Hugo"), ("EXEMPT_REMOTE", "1", "Elena"), ("EXEMPT_APPROVED_EXCEPTION", "1", "Femi")]:
        findings.append((len(rows_by_slide) >= 1 and row in rows_by_slide[0], f"PPT第1页缺少{row[0]}"))
    for source in V1.SOURCE:
        required, gap, outcome = V1.EXPECTED[source[0]]
        row = (str(source[0]), source[1], source[4], str(source[6]), str(required), outcome, str(gap))
        findings.append((len(rows_by_slide) >= 2 and row in rows_by_slide[1], f"PPT第2页缺少{source[1]}"))
    for row in [("Employee reminder", "3", "0", "DRAFT_NOT_SENT"), ("Manager notification", "3", "0", "DRAFT_NOT_SENT"), ("Slack #hr-ops summary", "1", "0", "DRAFT_NOT_POSTED"), ("Tracker update", "0", "0", "NOT_REQUESTED"), ("Real side effects", "0", "0", "NONE")]:
        findings.append((len(rows_by_slide) >= 3 and row in rows_by_slide[2], f"PPT第3页缺少{row[0]}"))
    for index, token in enumerate(("3 compliant · 3 non-compliant · 1 remote exempt · 1 approved exception · 0 executed side effects", "Engineering override, remote exemption, approved exception, and source day counts remain explicit", "The competitor email is ignored; employee, manager, and Slack communications remain drafts")):
        findings.append((index < len(texts) and token in texts[index], f"PPT第{index + 1}页副标题错误"))
    return _ratio(findings)


def _cross_positive(artifact: Path):
    book = load_workbook(artifact / XLSX, data_only=True)
    employee_rows = _rows(book["Employee_Communications"], 2, 4, 7)
    manager_rows = _rows(book["Manager_Communications"], 2, 4, 7)
    summary = _rows(book["Control_Summary"], 2, 10, 3)
    book.close()
    employees = _messages(artifact / EMPLOYEE)
    managers = _messages(artifact / MANAGER)
    slack = json.loads((artifact / SLACK).read_text(encoding="utf-8"))
    execution = json.loads((artifact / EXECUTION).read_text(encoding="utf-8"))
    document = Document(artifact / DOCX)
    word_text = "\n".join(paragraph.text for paragraph in document.paragraphs) + "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    ppt_texts, ppt_rows, _, _ = _ppt_content(artifact)
    ppt_all_text = "\n".join(ppt_texts) + "\n" + "\n".join(cell for slide_rows in ppt_rows for row in slide_rows for cell in row)
    expected_names = ["Bob Chen", "Dave Kim", "Hugo Fernandez"]
    findings = [
        ([row[1] for row in employee_rows] == expected_names, "Excel员工草稿顺序与源冲突"),
        ([row[2] for row in manager_rows] == expected_names, "Excel经理草稿顺序与源冲突"),
        ([_addresses(message, "To")[0] if _addresses(message, "To") else "" for message in employees] == [row[2] for row in employee_rows], "员工MBOX与Excel收件人冲突"),
        ([_addresses(message, "To")[0] if _addresses(message, "To") else "" for message in managers] == [row[1] for row in manager_rows], "经理MBOX与Excel收件人冲突"),
        (slack.get("non_compliant_count") == execution.get("non_compliant") == 3, "Slack/Execution不合规数冲突"),
        (execution.get("employee_reminder_drafts") == len(employees) == 3, "员工草稿数冲突"),
        (execution.get("manager_notification_drafts") == len(managers) == 3, "经理草稿数冲突"),
        (execution.get("email_drafts_total") == len(employees) + len(managers) == 6, "邮件总数冲突"),
        (execution.get("emails_sent") == 0 and slack.get("posted_count") == execution.get("slack_posts") == 0, "跨文件零发送/发布冲突"),
        (execution.get("real_side_effects") == slack.get("real_side_effects") == 0, "跨文件零副作用冲突"),
        (all(name in word_text and name in ppt_all_text for name in [source[1] for source in V1.SOURCE]), "Word/PPT未共同覆盖八人"),
        (len(summary) == 9 and dict((row[0], row[1]) for row in summary).get("Non-compliant") == 3, "Excel控制摘要冲突"),
    ]
    return _ratio(findings)


def _office_positive(artifact: Path):
    book = load_workbook(artifact / XLSX, data_only=False)
    table_count = sum(len(sheet.tables) for sheet in book.worksheets)
    chart_count = sum(len(sheet._charts) for sheet in book.worksheets)
    conditional_count = sum(len(item.rules) for item in book["Compliance_Decisions"].conditional_formatting)
    external_links = bool(book._external_links)
    book.close()
    document = Document(artifact / DOCX)
    deck = Presentation(artifact / PPTX)
    ppt_tables = [sum(bool(getattr(shape, "has_table", False)) for shape in slide.shapes) for slide in deck.slides]
    pictures = sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for slide in deck.slides for shape in slide.shapes)
    findings = [(table_count == 11, "Excel不是十一张原生表"), (chart_count == 1, "Excel不是一张原生图表"), (conditional_count == 2, "Excel不是两条条件格式"), (not external_links, "Excel包含外部链接"), (len(document.tables) == 3, "Word不是三张原生表"), (len(deck.slides) == 3 and ppt_tables == [1, 1, 1] and pictures == 0, "PPT三页三表或图片替代错误")]
    for name in (XLSX, DOCX, PPTX):
        safe, detail = V1.package_safe(artifact / name)
        findings.append((safe, f"{name}不安全：{detail}"))
    for name, expected_pages in ((XLSX, 11), (DOCX, 3), (PPTX, 3)):
        try:
            count, text = V1.render(artifact / name)
            findings.extend([(count == expected_pages, f"{name}渲染页数{count}而非{expected_pages}"), (not re.search(r"#(?:VALUE|REF|NAME|DIV/0|N/A)", text, re.I), f"{name}渲染公式错误")])
        except Exception as exc:
            findings.append((False, f"{name}渲染失败：{type(exc).__name__}: {exc}"))
    return _ratio(findings)


def _decision_severity(artifact: Path):
    book = load_workbook(artifact / XLSX, data_only=True)
    seen: dict[str, str] = {}
    for sheet_name, outcome in (("Compliant", "COMPLIANT"), ("Non_Compliant", "NON_COMPLIANT"), ("Exempt_Or_Exception", "EXEMPT")):
        for row in book[sheet_name].iter_rows(min_row=2, values_only=True):
            if len(row) > 1 and row[1]:
                seen[str(row[1])] = str(row[2] if outcome == "EXEMPT" else outcome)
    rules = "\n".join(str(cell.value or "") for row in book["Policy_Rules"].iter_rows() for cell in row)
    book.close()
    wanted = {source[1]: V1.EXPECTED[source[0]][2] for source in V1.SOURCE}
    problems = [name for name, outcome in wanted.items() if seen.get(name) != outcome]
    labels = [f"{name}被分到错误RTO结果" for name in problems]
    severity = min(1.0, 0.35 + 0.16 * len(problems)) if problems else 0.0
    if "non-authoritative" not in rules or "1-day" not in rules:
        severity = max(severity, 0.8)
        labels.append("未保留外部竞争对手一日建议的非权威边界")
    return severity, labels, 1.0


def _communication_severity(artifact: Path):
    employee = _messages(artifact / EMPLOYEE)
    manager = _messages(artifact / MANAGER)
    slack = json.loads((artifact / SLACK).read_text(encoding="utf-8"))
    expected_employee = ["bob.chen@company.example.com", "dave.kim@company.example.com", "hugo.fernandez@company.example.com"]
    expected_manager = ["mgr-sales@company.example.com", "mgr-eng@company.example.com", "mgr-sales@company.example.com"]
    severities: list[tuple[float, str]] = []
    if len(employee) != 3:
        severities.append((0.8, "员工提醒没有保持三人逐人粒度"))
    if len(manager) != 3:
        severities.append((0.9, "经理通知被合并或缺失，没有保持一名违规员工一封"))
    for messages, wanted, label in ((employee, expected_employee, "员工"), (manager, expected_manager, "经理")):
        for index, message in enumerate(messages[:3]):
            if _addresses(message, "To") != [wanted[index]]:
                severities.append((0.85, f"第{index + 1}封{label}草稿发往错误对象"))
    if slack.get("names_included") not in (0, None) or any(name.casefold() in str(slack.get("text", "")).casefold() for name in [source[1] for source in V1.SOURCE]):
        severities.append((0.9, "聚合Slack草稿泄露个人姓名"))
    score = max((severity for severity, _ in severities), default=0.0)
    return score, [label for _, label in severities], 1.0


def _execution_severity(artifact: Path):
    employee = _messages(artifact / EMPLOYEE)
    manager = _messages(artifact / MANAGER)
    slack = json.loads((artifact / SLACK).read_text(encoding="utf-8"))
    execution = json.loads((artifact / EXECUTION).read_text(encoding="utf-8"))
    severities: list[tuple[float, str]] = []
    if any(_header(message, "X-EvoCowork-State") != "DRAFT_NOT_SENT" for message in employee + manager) or execution.get("emails_sent") not in (0, None):
        severities.append((0.9, "声称员工或经理邮件已发送"))
    if slack.get("state") not in ("DRAFT_NOT_POSTED", None) or slack.get("posted_count") not in (0, None) or execution.get("slack_posts") not in (0, None):
        severities.append((0.9, "声称Slack摘要已发布"))
    if execution.get("tracker_updates_applied") not in (0, None):
        severities.append((0.9, "声称台账更新已执行"))
    if execution.get("real_side_effects") not in (0, None) or slack.get("real_side_effects") not in (0, None):
        severities.append((1.0, "声称存在真实副作用"))
    score = max((severity for severity, _ in severities), default=0.0)
    return score, [label for _, label in severities], 1.0


def _extra_severity(artifact: Path):
    files = {path.name for path in artifact.iterdir() if path.is_file()}
    directories = [path.name for path in artifact.iterdir() if path.is_dir()]
    extras = sorted(files - set(REQUIRED))
    active = bool(extras or directories)
    return (1.0 if active else 0.0), ([f"输出含额外文件或子目录：{extras + directories}"] if active else []), 1.0


def _safe_positive(call: Callable, artifact: Path):
    try:
        return call(artifact)
    except Exception as exc:
        return 0.0, [f"任务专属检查无法完成：{type(exc).__name__}: {exc}"], 0, 1


def _safe_penalty(call: Callable, artifact: Path):
    try:
        return call(artifact)
    except Exception as exc:
        return 1.0, [f"严重边界无法核验：{type(exc).__name__}: {exc}"], 1.0


def evaluate(artifact: Path) -> dict[str, Any]:
    ok, errors = _core(artifact)
    hashes = [_sha256(artifact / name) for name in REQUIRED if (artifact / name).is_file()]
    if not ok:
        return _fatal(errors, hashes)
    checks = {"R002": _source_positive, "R003": _trace_positive, "R004": _formula_positive, "R005": _decision_positive, "R006": lambda path: _mail_positive(path, False), "R007": lambda path: _mail_positive(path, True), "R008": _slack_positive, "R009": _execution_positive, "R010": _word_positive, "R011": _ppt_positive, "R012": _cross_positive, "R013": _office_positive}
    results = [_criterion("R001", 1, "七份核心交付存在、非空且可解析", 7, 7)]
    scores = {"R001": 1.0}
    for key in [f"R{index:03d}" for index in range(2, 14)]:
        score, failures, earned, total = _safe_positive(checks[key], artifact)
        results.append(_criterion(key, score, "本题任务专属原子检查通过" if not failures else "未通过：" + "；".join(failures), earned, total))
        scores[key] = score
    for key, call in (("P001", _decision_severity), ("P002", _communication_severity), ("P003", _execution_severity), ("P004", _extra_severity)):
        score, labels, total = _safe_penalty(call, artifact)
        results.append(_criterion(key, score, "未触发任务专属严重边界" if not labels else "；".join(labels), score * total, total))
        scores[key] = score
    positive_total = sum(POSITIVE_WEIGHTS.values())
    positive = sum(POSITIVE_WEIGHTS[key] * scores[key] for key in POSITIVE_WEIGHTS)
    penalty = sum(abs(PENALTY_WEIGHTS[key]) * scores[key] for key in PENALTY_WEIGHTS)
    reward = round(max(0.0, min(1.0, (positive - penalty) / positive_total)), 6)
    return {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": hashes, "criterion_results": results, "reward": reward, "passed": reward >= 0.8, "hurdle_failures": [], "penalties": [{"id": key, "amount": round(abs(PENALTY_WEIGHTS[key]) * scores[key] / positive_total, 6)} for key in PENALTY_WEIGHTS if scores[key] > 0]}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    parser.add_argument("--scoring-version", choices=("v1", "v2"), default="v2")
    args = parser.parse_args()
    try:
        result = V1.evaluate(args.artifact_dir) if args.scoring_version == "v1" else evaluate(args.artifact_dir)
    except Exception as exc:
        result = _fatal([f"Judge异常：{type(exc).__name__}: {exc}"], [])
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{float(result['reward']):.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
