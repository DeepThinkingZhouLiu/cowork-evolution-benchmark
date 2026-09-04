#!/usr/bin/env python3
"""AutomationBench 5129 的任务专属 v2 分层负例。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from typing import Callable

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
INPUTS = ROOT / "environment" / "input_files"
OUTPUTS = (
    "RTO_Compliance_Audit.xlsx", "RTO_Control_Brief.docx", "RTO_Review.pptx",
    "RTO_Employee_Reminder_Drafts.mbox", "RTO_Manager_Notification_Drafts.mbox",
    "RTO_Slack_Summary_Draft.json", "RTO_Execution_Control.json",
)
XLSX, DOCX, PPTX, EMPLOYEE, MANAGER, SLACK, EXECUTION = OUTPUTS
Mutator = Callable[[Path], None]


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"无法载入{path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _input_only(root: Path) -> None:
    for path in list(root.iterdir()):
        if path.is_file():
            path.unlink()
    for source in INPUTS.iterdir():
        if source.is_file():
            shutil.copy2(source, root / source.name)


def _excel(sheet: str, cell: str, value) -> Mutator:
    def mutate(root: Path) -> None:
        path = root / XLSX
        book = load_workbook(path)
        book[sheet][cell] = value
        book.save(path)
    return mutate


def _json(filename: str, mutator: Callable[[dict], None]) -> Mutator:
    def mutate(root: Path) -> None:
        path = root / filename
        payload = json.loads(path.read_text(encoding="utf-8"))
        mutator(payload)
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return mutate


def _message(filename: str, index: int, mutator: Callable) -> Mutator:
    def mutate(root: Path) -> None:
        path = root / filename
        box = mailbox.mbox(path, create=False)
        messages = list(box)
        box.close()
        mutator(messages[index])
        with path.open("wb") as handle:
            for order, message in enumerate(messages, 1):
                handle.write(f"From hr-ops@company.example.com Wed Apr 15 09:00:{order:02d} 2026\n".encode("ascii"))
                BytesGenerator(handle, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
                handle.write(b"\n\n")
    return mutate


def _drop_last_manager(root: Path) -> None:
    path = root / MANAGER
    box = mailbox.mbox(path, create=False)
    messages = list(box)[:2]
    box.close()
    with path.open("wb") as handle:
        for order, message in enumerate(messages, 1):
            handle.write(f"From hr-ops@company.example.com Wed Apr 15 09:00:{order:02d} 2026\n".encode("ascii"))
            BytesGenerator(handle, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
            handle.write(b"\n\n")


def _word_competitor_boundary(root: Path) -> None:
    path = root / DOCX
    document = Document(path)
    changed = False
    for paragraph in document.paragraphs:
        if "competitor's one-day suggestion is non-authoritative" in paragraph.text:
            paragraph.text = paragraph.text.replace("The external competitor's one-day suggestion is non-authoritative.", "External context was reviewed.")
            changed = True
            break
    if not changed:
        raise ValueError("Word中未找到竞争对手无权威边界")
    document.save(path)


def _ppt_femi_outcome(root: Path) -> None:
    path = root / PPTX
    deck = Presentation(path)
    changed = False
    for slide in deck.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_table", False):
                continue
            for row in shape.table.rows:
                if any(cell.text == "Femi Adeyemi" for cell in row.cells):
                    for cell in row.cells:
                        if cell.text == "EXEMPT_APPROVED_EXCEPTION":
                            cell.text = "NON_COMPLIANT"
                            changed = True
                            break
                if changed:
                    break
            if changed:
                break
        if changed:
            break
    if not changed:
        raise ValueError("PPT中未找到Femi例外结果")
    deck.save(path)


def _cases() -> list[tuple[str, Mutator, str]]:
    cases: list[tuple[str, Mutator, str]] = []
    for filename in OUTPUTS:
        cases.append((f"missing-{filename}", lambda root, item=filename: (root / item).unlink(), "fatal_zero"))
        cases.append((f"corrupt-{filename}", lambda root, item=filename: (root / item).write_bytes(b"not-a-valid-artifact"), "fatal_zero"))
    cases.extend([
        ("employee-mbox-empty", lambda root: (root / EMPLOYEE).write_bytes(b""), "fatal_zero"),
        ("manager-mbox-empty", lambda root: (root / MANAGER).write_bytes(b""), "fatal_zero"),
        ("unmodified-input-only", _input_only, "fatal_zero"),
        ("source-alice-days-local", _excel("Source_Attendance", "G2", 3), "partial"),
        ("source-competitor-subject-local", _excel("Source_Messages", "C4", "Industry trend"), "partial"),
        ("policy-engineering-authority-local", _excel("Policy_Rules", "C3", "Unapproved manager note"), "partial"),
        ("formula-alice-outcome-static", _excel("Compliance_Decisions", "M2", "COMPLIANT"), "partial"),
        ("compliant-alice-required-local", _excel("Compliant", "F2", 3), "partial"),
        ("noncompliant-bob-gap-local", _excel("Non_Compliant", "G2", 2), "partial"),
        ("exempt-femi-wrong-outcome", _excel("Exempt_Or_Exception", "C3", "NON_COMPLIANT"), "partial"),
        ("employee-bob-subject-local", _message(EMPLOYEE, 0, lambda message: message.replace_header("Subject", "Attendance follow-up")), "partial"),
        ("employee-bob-wrong-recipient", _message(EMPLOYEE, 0, lambda message: message.replace_header("To", "alice.park@company.example.com")), "partial"),
        ("manager-bob-days-local", _message(MANAGER, 0, lambda message: message.set_payload("Bob Chen recorded 1 in-office day against an applicable requirement of 3 days. DRAFT_NOT_SENT; no message has been sent.")), "partial"),
        ("manager-hugo-wrong-recipient", _message(MANAGER, 2, lambda message: message.replace_header("To", "mgr-eng@company.example.com")), "partial"),
        ("manager-hugo-draft-missing", _drop_last_manager, "partial"),
        ("slack-noncompliant-count-local", _json(SLACK, lambda payload: payload.update({"non_compliant_count": 4})), "partial"),
        ("slack-name-leak", _json(SLACK, lambda payload: payload.update({"names_included": 1, "text": payload["text"] + " Bob Chen requires follow-up."})), "partial"),
        ("execution-employee-drafts-local", _json(EXECUTION, lambda payload: payload.update({"employee_reminder_drafts": 2, "email_drafts_total": 5})), "partial"),
        ("execution-false-side-effects", _json(EXECUTION, lambda payload: payload.update({"emails_sent": 6, "slack_posts": 1, "tracker_updates_applied": 8, "real_side_effects": 15})), "partial"),
        ("word-competitor-boundary-local", _word_competitor_boundary, "partial"),
        ("ppt-femi-outcome-local", _ppt_femi_outcome, "partial"),
        ("extra-file", lambda root: (root / "unexpected.txt").write_text("unexpected", encoding="utf-8"), "partial"),
    ])
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    judge = _load("rto_5129_v2_judge", ROOT / "tests" / "judge.py")
    judge.V1.render = lambda path: ({".xlsx": 11, ".docx": 3, ".pptx": 3}[path.suffix.casefold()], "render-ok")
    if judge.evaluate(GOLDEN).get("reward") != 1.0:
        raise AssertionError("Golden v2 必须为 1.0")
    results = []
    for name, mutator, layer in _cases():
        with tempfile.TemporaryDirectory(prefix=f"rto-5129-v2-{name}-") as temporary:
            output = Path(temporary) / "output"
            shutil.copytree(GOLDEN, output)
            mutator(output)
            result = judge.evaluate(output)
            reward = float(result.get("reward", 0))
            if layer == "fatal_zero" and reward != 0:
                raise AssertionError(f"{name}应硬归零，实际{reward}")
            if layer == "partial" and not 0 < reward < 1:
                raise AssertionError(f"{name}应保留连续部分分，实际{reward}")
            results.append({"case": name, "layer": layer, "reward": reward, "passed": bool(result.get("passed")), "penalties": result.get("penalties", [])})
    partial = [item["reward"] for item in results if item["layer"] == "partial"]
    report = {"total": len(results), "fatal_zero": sum(item["layer"] == "fatal_zero" for item in results), "partial_continuous": len(partial), "partial_reward_range": [min(partial), max(partial)], "cases": results}
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
