#!/usr/bin/env python3
"""AutomationBench 5129 的任务定向负例回放。"""

from __future__ import annotations

import json, mailbox, shutil, subprocess, tempfile
import sys
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]; GOLDEN=ROOT/"solution"/"Golden"; PYTHON=sys.executable

def judge(case):
    result=case.parent/f"{case.name}-result.json";reward=case.parent/f"{case.name}-reward.txt";subprocess.run([PYTHON,str(ROOT/"tests"/"judge.py"),"--artifact-dir",str(case),"--result",str(result),"--reward-file",str(reward)],check=True,capture_output=True,text=True,timeout=240);return float(reward.read_text())

def retarget(path):
    box=mailbox.mbox(path,create=False);messages=[m for m in box];box.close();messages[0].replace_header("To","alice.park@company.example.com")
    with path.open("wb") as h:
        for i,m in enumerate(messages):h.write(f"From hr-ops@company.example.com Wed Apr 15 09:00:{i:02d} 2026\n".encode());BytesGenerator(h,policy=SMTP,mangle_from_=True,maxheaderlen=0).flatten(m);h.write(b"\n\n")

def main():
    with tempfile.TemporaryDirectory(prefix="rto-5129-negative-") as temp:
        root=Path(temp);names=["missing","wrong_counts","slack_posted","wrong_employee","wrong_manager","formula_removed","word_false","ppt_omission"];cases=[]
        for name in names:case=root/name;shutil.copytree(GOLDEN,case);cases.append(case)
        (cases[0]/"RTO_Review.pptx").unlink()
        p=json.loads((cases[1]/"RTO_Execution_Control.json").read_text());p["non_compliant"]=8;p["compliant"]=0;(cases[1]/"RTO_Execution_Control.json").write_text(json.dumps(p))
        p=json.loads((cases[2]/"RTO_Slack_Summary_Draft.json").read_text());p["posted_count"]=1;p["state"]="POSTED";p["real_side_effects"]=1;(cases[2]/"RTO_Slack_Summary_Draft.json").write_text(json.dumps(p))
        retarget(cases[3]/"RTO_Employee_Reminder_Drafts.mbox")
        retarget(cases[4]/"RTO_Manager_Notification_Drafts.mbox")
        wb=load_workbook(cases[5]/"RTO_Compliance_Audit.xlsx");wb["Compliance_Decisions"]["M2"]="NON_COMPLIANT";wb.save(cases[5]/"RTO_Compliance_Audit.xlsx")
        doc=Document(cases[6]/"RTO_Control_Brief.docx");doc.add_paragraph("All employees are non-compliant and messages were sent.");doc.save(cases[6]/"RTO_Control_Brief.docx")
        prs=Presentation(cases[7]/"RTO_Review.pptx")
        for slide in prs.slides:
            for shape in slide.shapes:
                if getattr(shape,"has_text_frame",False):shape.text=shape.text.replace("Femi Adeyemi","Unknown")
                if getattr(shape,"has_table",False):
                    for row in shape.table.rows:
                        for cell in row.cells:cell.text=cell.text.replace("Femi Adeyemi","Unknown")
        prs.save(cases[7]/"RTO_Review.pptx")
        rewards={case.name:judge(case) for case in cases}
        if any(v!=0 for v in rewards.values()):raise SystemExit(f"负例未归零:{rewards}")
        print(json.dumps({"total":len(rewards),"rejected":len(rewards),"maximum_reward":max(rewards.values()),"rewards":rewards},ensure_ascii=False,indent=2))

if __name__=="__main__":main()
