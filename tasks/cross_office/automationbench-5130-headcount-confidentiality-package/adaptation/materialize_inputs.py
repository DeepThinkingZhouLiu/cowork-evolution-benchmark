#!/usr/bin/env python3
"""只静态化 AutomationBench 5130 的真实初始状态；不是通用造题器。"""
from __future__ import annotations
import hashlib, importlib.util, json, shutil, sys, types
from datetime import datetime
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.worksheet.table import Table,TableStyleInfo

ROOT=Path(__file__).resolve().parents[1];UPSTREAM=Path(__file__).with_name("upstream-hr-tasks.py");SNAPSHOT=Path(__file__).with_name("upstream-task-5130.json");DATA=ROOT/"data"/"input_files";ENV=ROOT/"environment"/"input_files";SOURCE_SHA="c173a36a03f43096bd122577584aa32f9df379d8760aa46c43be04c85d101779"
def sha(path):
 d=hashlib.sha256();d.update(path.read_bytes());return d.hexdigest()
def load_task():
 if sha(UPSTREAM)!=SOURCE_SHA:raise ValueError("冻结上游哈希错误")
 datasets=types.ModuleType("datasets");datasets.Dataset=object;sys.modules["datasets"]=datasets
 for name in ("automationbench","automationbench.domains","automationbench.domains.hr"):sys.modules[name]=types.ModuleType(name)
 noise=types.ModuleType("automationbench.domains.hr._noise");noise.apply_noise=lambda x:x;sys.modules[noise.__name__]=noise
 spec=importlib.util.spec_from_file_location("ab5130",UPSTREAM);module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module);task=module.get_task_5130()
 if task["example_id"]!=5130 or task["task"]!="hr.headcount_planning":raise ValueError("上游任务错误")
 return task
def assert_shape(task):
 state=task["info"]["initial_state"];assert state["meta"]["current_time"]=="2026-04-15T09:00:00Z";assert len(task["info"]["assertions"])==10
 prompt=task["prompt"][1]["content"];assert all(x in prompt for x in ("Q2 headcount summary","Send it to the CEO","leadership channel","#all-staff"))
 msg=state["gmail"]["messages"];assert len(msg)==1 and msg[0]["id"]=="msg_headcount_approval_5130";body=msg[0]["body_plain"]
 assert all(x in body for x in ("Engineering: 2 additional","total approved goes from 15 to 17","Sales: 1 additional","CONFIDENTIAL pre-board-meeting","Do NOT post","#all-staff"))
 rows=state["google_sheets"]["spreadsheets"][0]["worksheets"][0]["rows"];assert [r["row_id"] for r in rows]==list(range(2,10));assert [r["cells"]["Department"] for r in rows]==["Engineering","Sales","Marketing","Support","Product","Legal","Finance","Design"]
 assert [r["cells"]["Approved Headcount"] for r in rows]==["15","10","8","6","5","3","4","3"];assert "Hiring freeze" in rows[-1]["cells"]["Notes"]
 assert state["slack"]["channels"]==[{"id":"C_LEAD_5130","name":"leadership","is_private":False},{"id":"C_ALL_5130","name":"all-staff","is_private":False}]
def message_bytes(record):
 m=EmailMessage(policy=SMTP);m["From"]=record["from_"];m["To"]=", ".join(record["to"]);m["Subject"]=record["subject"];parsed=datetime.fromisoformat(record["date"].replace("Z","+00:00"));m["Date"]=parsed.strftime("%a, %d %b %Y %H:%M:%S +0000");m["Message-ID"]=f"<{record['id']}@automationbench.invalid>";m["X-AutomationBench-Message-ID"]=record["id"];m.set_content(record["body_plain"],charset="utf-8");b=BytesIO();BytesGenerator(b,policy=SMTP,mangle_from_=True,maxheaderlen=0).flatten(m);return b.getvalue()
def write_mbox(messages,target):
 chunks=[]
 for r in messages:
  parsed=datetime.fromisoformat(r["date"].replace("Z","+00:00"));chunks += [f"From {r['from_']} {parsed.strftime('%a %b %d %H:%M:%S %Y')}\n".encode(),message_bytes(r).replace(b"\r\n",b"\n"),b"\n"]
 target.write_bytes(b"".join(chunks))
def write_tracker(state,target):
 source=state["google_sheets"]["spreadsheets"][0];rows=source["worksheets"][0]["rows"];wb=Workbook();ws=wb.active;ws.title="Department Summary";headers=["Source Row","Department","Approved Headcount","Filled","Open","Notes"];ws.append(headers)
 for item in rows:
  c=item["cells"];ws.append([item["row_id"],c["Department"],int(c["Approved Headcount"]),int(c["Filled"]),int(c["Open"]),c["Notes"]])
 ws.freeze_panes="A2";ws.sheet_view.showGridLines=False
 for i,w in enumerate([11,22,22,13,13,55],1):ws.column_dimensions[ws.cell(1,i).column_letter].width=w
 for c in ws[1]:c.fill=PatternFill("solid",fgColor="17365D");c.font=Font(name="Carlito",bold=True,color="FFFFFF");c.alignment=Alignment(horizontal="center",wrap_text=True)
 for row in ws.iter_rows(min_row=2):
  for c in row:c.font=Font(name="Carlito",size=9);c.alignment=Alignment(vertical="top",wrap_text=True)
 t=Table(displayName="Q2HeadcountSource",ref=ws.dimensions);t.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True);ws.add_table(t);ws.page_setup.orientation="landscape";ws.page_setup.fitToWidth=1;ws.sheet_properties.pageSetUpPr.fitToPage=True;wb.save(target)
def main():
 task=load_task();assert_shape(task);state=task["info"]["initial_state"];DATA.mkdir(parents=True,exist_ok=True);ENV.mkdir(parents=True,exist_ok=True);snap=json.dumps(task,ensure_ascii=False,indent=2)+"\n"
 if SNAPSHOT.exists() and SNAPSHOT.read_text()!=snap:raise ValueError("快照漂移")
 SNAPSHOT.write_text(snap);write_tracker(state,DATA/"Q2_Headcount_Plan.xlsx");write_mbox(state["gmail"]["messages"],DATA/"Headcount_Approval_and_Policy.mbox");(DATA/"Headcount_Slack_Channels.json").write_text(json.dumps(state["slack"],ensure_ascii=False,indent=2)+"\n");(DATA/"Workflow_Context.json").write_text(json.dumps({"meta":state["meta"]},indent=2)+"\n")
 for p in DATA.iterdir():
  if p.is_file():shutil.copy2(p,ENV/p.name)
if __name__=="__main__":main()
