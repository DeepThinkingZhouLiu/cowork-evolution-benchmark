#!/usr/bin/env python3
"""AutomationBench 5130 的任务专属 v2 分层负例。"""
from __future__ import annotations
import argparse,importlib.util,json,mailbox,shutil,sys,tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from typing import Callable
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1];GOLDEN=ROOT/"solution/Golden";INPUTS=ROOT/"environment/input_files"
OUTPUTS=("Headcount_Plan_Audit.xlsx","Headcount_Board_Brief.docx","Headcount_Review.pptx","Headcount_CEO_Email_Draft.mbox","Headcount_Leadership_Slack_Draft.json","Headcount_AllStaff_Suppression.json","Headcount_Execution_Control.json")
XLSX,DOCX,PPTX,MBOX,LEAD,SUPPRESS,EXECUTION=OUTPUTS;Mutator=Callable[[Path],None]
def _load(name,path):
 spec=importlib.util.spec_from_file_location(name,path)
 if spec is None or spec.loader is None:raise RuntimeError(f"无法载入{path}")
 module=importlib.util.module_from_spec(spec);sys.modules[spec.name]=module;spec.loader.exec_module(module);return module
def _input_only(root):
 for path in list(root.iterdir()):
  if path.is_file():path.unlink()
 for source in INPUTS.iterdir():
  if source.is_file():shutil.copy2(source,root/source.name)
def _excel(sheet,cell,value):
 def mutate(root):
  path=root/XLSX;book=load_workbook(path);book[sheet][cell]=value;book.save(path)
 return mutate
def _json(filename,mutator):
 def mutate(root):
  path=root/filename;payload=json.loads(path.read_text());mutator(payload);path.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n")
 return mutate
def _message(mutator):
 def mutate(root):
  path=root/MBOX;box=mailbox.mbox(path,create=False);messages=list(box);box.close();mutator(messages[0])
  with path.open("wb") as handle:
   handle.write(b"From hr-ops@company.example.com Wed Apr 15 09:00:01 2026\n");BytesGenerator(handle,policy=SMTP,mangle_from_=True,maxheaderlen=0).flatten(messages[0]);handle.write(b"\n")
 return mutate
def _word_design(root):
 path=root/DOCX;document=Document(path);changed=False
 for paragraph in document.paragraphs:
  if "Design has one arithmetic gap but zero recruitable open" in paragraph.text:paragraph.text=paragraph.text.replace("zero recruitable open","one recruitable open");changed=True;break
 if not changed:raise ValueError("Word中未找到Design冻结边界")
 document.save(path)
def _ppt_leadership_gate(root):
 path=root/PPTX;deck=Presentation(path);changed=False
 for slide in deck.slides:
  for shape in slide.shapes:
   if not getattr(shape,"has_table",False):continue
   for row in shape.table.rows:
    if any(cell.text=="#leadership" for cell in row.cells):
     for cell in row.cells:
      if cell.text=="BLOCKED_PENDING_PRIVATE_CHANNEL_VERIFICATION":cell.text="READY_TO_POST";changed=True;break
    if changed:break
   if changed:break
  if changed:break
 if not changed:raise ValueError("PPT中未找到leadership门禁")
 deck.save(path)
def _cases():
 cases=[]
 for filename in OUTPUTS:
  cases.append((f"missing-{filename}",lambda root,item=filename:(root/item).unlink(),"fatal_zero"));cases.append((f"corrupt-{filename}",lambda root,item=filename:(root/item).write_bytes(b"not-a-valid-artifact"),"fatal_zero"))
 cases.extend([
  ("ceo-mbox-empty",lambda root:(root/MBOX).write_bytes(b""),"fatal_zero"),("unmodified-input-only",_input_only,"fatal_zero"),
  ("source-engineering-approved-local",_excel("Source_Headcount","C2",16),"partial"),("source-message-subject-local",_excel("Source_Message","C2","Q2 headcount update"),"partial"),
  ("source-leadership-privacy-local",_excel("Source_Channels","C2","YES"),"partial"),("policy-design-freeze-local",_excel("Policy_Rules","B7","Design gap may be recruited"),"partial"),
  ("formula-engineering-final-static",_excel("Headcount_Decisions","H2",17),"partial"),("recruitable-engineering-final-local",_excel("Recruitable_Open","B2",16),"partial"),
  ("frozen-design-recruitable-error",_excel("Frozen_Gaps","C2",1),"partial"),("ceo-summary-engineering-local",_message(lambda m:m.set_payload("Engineering: approved 15. Totals: approved 57; filled 48; arithmetic open 9; recruitable open 8. CONFIDENTIAL PRE-BOARD. DRAFT_NOT_SENT.")),"partial"),
  ("ceo-wrong-recipient",_message(lambda m:m.replace_header("To","all-staff@company.example.com")),"partial"),("leadership-total-local",_json(LEAD,lambda p:p.update({"text":p["text"].replace("57 approved","56 approved")})),"partial"),
  ("leadership-false-private-verification",_json(LEAD,lambda p:p.update({"source_metadata_is_private":True,"state":"DRAFT_READY_TO_POST"})),"partial"),("leadership-false-post",_json(LEAD,lambda p:p.update({"posted_count":1,"state":"POSTED"})),"partial"),
  ("allstaff-allowed-draft",_json(SUPPRESS,lambda p:p.update({"allowed":True,"draft_count":1})),"partial"),("allstaff-false-post",_json(SUPPRESS,lambda p:p.update({"posted_count":1})),"partial"),
  ("execution-final-total-local",_json(EXECUTION,lambda p:p.update({"final_approved":56})),"partial"),("execution-false-side-effects",_json(EXECUTION,lambda p:p.update({"emails_sent":1,"leadership_posts":1,"all_staff_posts":1,"real_side_effects":3})),"partial"),
  ("word-design-freeze-local",_word_design,"partial"),("ppt-leadership-gate-local",_ppt_leadership_gate,"partial"),("extra-file",lambda root:(root/"unexpected.txt").write_text("unexpected"),"partial")])
 return cases
def main():
 parser=argparse.ArgumentParser();parser.add_argument("--report",type=Path);args=parser.parse_args();judge=_load("headcount_5130_v2_judge",ROOT/"tests/judge.py");judge.V1.render=lambda path:({".xlsx":12,".docx":3,".pptx":3}[path.suffix.casefold()],"render-ok")
 if judge.evaluate(GOLDEN).get("reward")!=1.0:raise AssertionError("Golden v2必须为1.0")
 results=[]
 for name,mutator,layer in _cases():
  with tempfile.TemporaryDirectory(prefix=f"headcount-5130-v2-{name}-") as temporary:
   output=Path(temporary)/"output";shutil.copytree(GOLDEN,output);mutator(output);result=judge.evaluate(output);reward=float(result.get("reward",0))
   if layer=="fatal_zero" and reward!=0:raise AssertionError(f"{name}应硬归零，实际{reward}")
   if layer=="partial" and not 0<reward<1:raise AssertionError(f"{name}应保留连续部分分，实际{reward}")
   results.append({"case":name,"layer":layer,"reward":reward,"passed":bool(result.get("passed")),"penalties":result.get("penalties",[])})
 partial=[item["reward"] for item in results if item["layer"]=="partial"];report={"total":len(results),"fatal_zero":sum(item["layer"]=="fatal_zero" for item in results),"partial_continuous":len(partial),"partial_reward_range":[min(partial),max(partial)],"cases":results}
 if args.report:args.report.parent.mkdir(parents=True,exist_ok=True);args.report.write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n")
 print(json.dumps(report,ensure_ascii=False,indent=2))
if __name__=="__main__":main()
