#!/usr/bin/env python3
"""AutomationBench 5130 的任务定向负例。"""
import json,mailbox,shutil,subprocess,tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1];G=ROOT/"solution"/"Golden";PY="REMOVED_LOCAL_PATH"
def judge(c):
 r=c.parent/(c.name+".json");w=c.parent/(c.name+".txt");subprocess.run([PY,str(ROOT/"tests/judge.py"),"--artifact-dir",str(c),"--result",str(r),"--reward-file",str(w)],check=True,capture_output=True,text=True,timeout=240);return float(w.read_text())
def mbox_to_allstaff(path):
 b=mailbox.mbox(path,create=False);ms=[m for m in b];b.close();ms[0].replace_header("To","all-staff@company.example.com")
 with path.open("wb") as h:
  h.write(b"From hr-ops@company.example.com Wed Apr 15 09:00:01 2026\n");BytesGenerator(h,policy=SMTP,mangle_from_=True,maxheaderlen=0).flatten(ms[0]);h.write(b"\n")
def main():
 with tempfile.TemporaryDirectory(prefix="hc5130-neg-") as t:
  root=Path(t);names=["missing","wrong_total","design_open","allstaff_allowed","leadership_posted","wrong_recipient","formula_removed","word_false","ppt_omit"];cases=[]
  for n in names:c=root/n;shutil.copytree(G,c);cases.append(c)
  (cases[0]/"Headcount_Review.pptx").unlink()
  p=json.loads((cases[1]/"Headcount_Execution_Control.json").read_text());p["final_approved"]=54;(cases[1]/"Headcount_Execution_Control.json").write_text(json.dumps(p))
  p=json.loads((cases[2]/"Headcount_Execution_Control.json").read_text());p["recruitable_open"]=9;(cases[2]/"Headcount_Execution_Control.json").write_text(json.dumps(p))
  p=json.loads((cases[3]/"Headcount_AllStaff_Suppression.json").read_text());p["allowed"]=True;p["posted_count"]=1;(cases[3]/"Headcount_AllStaff_Suppression.json").write_text(json.dumps(p))
  p=json.loads((cases[4]/"Headcount_Leadership_Slack_Draft.json").read_text());p["posted_count"]=1;p["state"]="POSTED";(cases[4]/"Headcount_Leadership_Slack_Draft.json").write_text(json.dumps(p))
  mbox_to_allstaff(cases[5]/"Headcount_CEO_Email_Draft.mbox")
  wb=load_workbook(cases[6]/"Headcount_Plan_Audit.xlsx");wb["Headcount_Decisions"]["H2"]=15;wb.save(cases[6]/"Headcount_Plan_Audit.xlsx")
  d=Document(cases[7]/"Headcount_Board_Brief.docx");d.add_paragraph("Posted to all-staff; channel privacy verified.");d.save(cases[7]/"Headcount_Board_Brief.docx")
  p=Presentation(cases[8]/"Headcount_Review.pptx")
  for s in p.slides:
   for sh in s.shapes:
    if getattr(sh,"has_text_frame",False):sh.text=sh.text.replace("Design","Unknown")
    if getattr(sh,"has_table",False):
     for r in sh.table.rows:
      for c in r.cells:c.text=c.text.replace("Design","Unknown")
  p.save(cases[8]/"Headcount_Review.pptx")
  rewards={c.name:judge(c) for c in cases};print(json.dumps(rewards,indent=2));assert all(v==0 for v in rewards.values()),rewards
if __name__=="__main__":main()
