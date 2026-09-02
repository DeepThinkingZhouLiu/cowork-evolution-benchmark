#!/usr/bin/env python3
"""只静态化 AutomationBench 5131；不是通用造题器。"""
import hashlib,importlib.util,json,shutil,sys,types
from datetime import datetime,date
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.worksheet.table import Table,TableStyleInfo
ROOT=Path(__file__).resolve().parents[1];UP=Path(__file__).with_name("upstream-hr-tasks.py");SNAP=Path(__file__).with_name("upstream-task-5131.json");DATA=ROOT/"data/input_files";ENV=ROOT/"environment/input_files";SHA="c173a36a03f43096bd122577584aa32f9df379d8760aa46c43be04c85d101779"
def load():
 if hashlib.sha256(UP.read_bytes()).hexdigest()!=SHA:raise ValueError("上游哈希错误")
 d=types.ModuleType("datasets");d.Dataset=object;sys.modules["datasets"]=d
 for n in ("automationbench","automationbench.domains","automationbench.domains.hr"):sys.modules[n]=types.ModuleType(n)
 no=types.ModuleType("automationbench.domains.hr._noise");no.apply_noise=lambda x:x;sys.modules[no.__name__]=no
 s=importlib.util.spec_from_file_location("ab5131",UP);m=importlib.util.module_from_spec(s);s.loader.exec_module(m);t=m.get_task_5131();assert t["example_id"]==5131 and t["task"]=="hr.handbook_acknowledgment";return t
def check(t):
 st=t["info"]["initial_state"];assert st["meta"]["current_time"]=="2026-04-15T09:00:00Z" and len(t["info"]["assertions"])==8
 msg=st["gmail"]["messages"];assert len(msg)==1 and msg[0]["id"]=="msg_handbook_compliance_5131";b=msg[0]["body_plain"]
 assert all(x in b for x in ("March 1, 2026","date before March 1","deadline: March 31, 2026","within 45 days (April 15)","approved leave","7 days after their return"))
 rs=st["google_sheets"]["spreadsheets"][0]["worksheets"][0]["rows"];assert [r["row_id"] for r in rs]==list(range(2,10));assert [r["cells"]["Employee"] for r in rs]==["Alice Park","Bob Chen","Carol Diaz","Dave Kim","Elena Vasquez","Femi Adeyemi","Grace Okonkwo","Hugo Fernandez"]
 assert rs[4]["cells"]["Notes"]=="On medical leave - returning April 20";assert st["slack"]["channels"]==[{"id":"C_COMP_5131","name":"compliance","is_private":False}]
def mb(messages,path):
 chunks=[]
 for r in messages:
  m=EmailMessage(policy=SMTP);m["From"]=r["from_"];m["To"]=", ".join(r["to"]);m["Subject"]=r["subject"];dt=datetime.fromisoformat(r["date"].replace("Z","+00:00"));m["Date"]=dt.strftime("%a, %d %b %Y %H:%M:%S +0000");m["Message-ID"]=f"<{r['id']}@automationbench.invalid>";m["X-AutomationBench-Message-ID"]=r["id"];m.set_content(r["body_plain"],charset="utf-8");buf=BytesIO();BytesGenerator(buf,policy=SMTP,mangle_from_=True,maxheaderlen=0).flatten(m);chunks += [f"From {r['from_']} {dt.strftime('%a %b %d %H:%M:%S %Y')}\n".encode(),buf.getvalue().replace(b"\r\n",b"\n"),b"\n"]
 path.write_bytes(b"".join(chunks))
def xlsx(st,path):
 rs=st["google_sheets"]["spreadsheets"][0]["worksheets"][0]["rows"];wb=Workbook();ws=wb.active;ws.title="Acknowledgments";hs=["Source Row","Employee","Email","Manager Email","Acknowledgment Date","Status","Notes"];ws.append(hs)
 for r in rs:
  c=r["cells"];ack=date.fromisoformat(c["Acknowledgment Date"]) if c["Acknowledgment Date"] else None;ws.append([r["row_id"],c["Employee"],c["Email"],c["Manager Email"],ack,c["Status"],c["Notes"]])
 for i,w in enumerate([11,22,32,32,22,22,55],1):ws.column_dimensions[ws.cell(1,i).column_letter].width=w
 ws.freeze_panes="A2";ws.sheet_view.showGridLines=False
 for c in ws[1]:c.fill=PatternFill("solid",fgColor="17365D");c.font=Font(name="Carlito",bold=True,color="FFFFFF");c.alignment=Alignment(horizontal="center",wrap_text=True)
 for r in ws.iter_rows(min_row=2):
  for c in r:c.font=Font(name="Carlito",size=9);c.alignment=Alignment(vertical="top",wrap_text=True)
  r[4].number_format="yyyy-mm-dd"
 t=Table(displayName="HandbookAcknowledgmentsSource",ref=ws.dimensions);t.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True);ws.add_table(t);ws.page_setup.orientation="landscape";ws.page_setup.fitToWidth=1;ws.sheet_properties.pageSetUpPr.fitToPage=True;wb.save(path)
def main():
 t=load();check(t);st=t["info"]["initial_state"];DATA.mkdir(parents=True,exist_ok=True);ENV.mkdir(parents=True,exist_ok=True);s=json.dumps(t,ensure_ascii=False,indent=2)+"\n"
 if SNAP.exists() and SNAP.read_text()!=s:raise ValueError("快照漂移")
 SNAP.write_text(s);xlsx(st,DATA/"Handbook_Acknowledgment_Tracker.xlsx");mb(st["gmail"]["messages"],DATA/"Handbook_Compliance_Policy.mbox");(DATA/"Compliance_Slack.json").write_text(json.dumps(st["slack"],ensure_ascii=False,indent=2)+"\n");(DATA/"Workflow_Context.json").write_text(json.dumps({"meta":st["meta"]},indent=2)+"\n")
 for p in DATA.iterdir():
  if p.is_file():shutil.copy2(p,ENV/p.name)
if __name__=="__main__":main()
