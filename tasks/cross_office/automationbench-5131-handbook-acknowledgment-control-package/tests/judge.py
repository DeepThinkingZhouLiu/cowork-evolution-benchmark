#!/usr/bin/env python3
"""AutomationBench 5131 的 v2 连续评分入口；原严格口径保存在 judge_v1.py。"""
from __future__ import annotations
import argparse,hashlib,importlib.util,json,mailbox,re,sys
from datetime import datetime
from email.header import decode_header,make_header
from email.utils import getaddresses
from pathlib import Path
from typing import Any
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
TASK_ID="cowork-evo/automationbench-5131-handbook-acknowledgment-control-package";RUBRIC_VERSION="automationbench-5131-handbook-v0.2";JUDGE_VERSION="automationbench-5131-criterion-partial-v2.0"
XLSX="Handbook_Compliance_Audit.xlsx";DOCX="Handbook_Control_Brief.docx";PPTX="Handbook_Review.pptx";EMPLOYEE="Handbook_Employee_Manager_Drafts.mbox";DIRECTOR="Handbook_HR_Director_Escalation_Draft.mbox";SLACK="Handbook_Slack_Summary_Draft.json";EXECUTION="Handbook_Execution_Control.json";REQUIRED=(XLSX,DOCX,PPTX,EMPLOYEE,DIRECTOR,SLACK,EXECUTION)
POSITIVE_WEIGHTS={"R001":5,"R002":8,"R003":7,"R004":12,"R005":12,"R006":8,"R007":10,"R008":8,"R009":7,"R010":7,"R011":7,"R012":4,"R013":5};PENALTY_WEIGHTS={"P001":-35,"P002":-30,"P003":-35,"P004":-10}
def _load_v1():
 path=Path(__file__).with_name("judge_v1.py")
 if str(path.parent) not in sys.path:sys.path.insert(0,str(path.parent))
 spec=importlib.util.spec_from_file_location("automationbench_5131_judge_v1",path)
 if spec is None or spec.loader is None:raise RuntimeError("无法载入v1 Judge")
 module=importlib.util.module_from_spec(spec);sys.modules[spec.name]=module;spec.loader.exec_module(module);return module
V1=_load_v1()
def _sha256(path):
 d=hashlib.sha256()
 with path.open("rb") as h:
  for chunk in iter(lambda:h.read(1048576),b""):d.update(chunk)
 return d.hexdigest()
def _rows(sheet,start,end,columns):return [tuple(sheet.cell(r,c).value for c in range(1,columns+1)) for r in range(start,end+1)]
def _messages(path):
 box=mailbox.mbox(path,create=False);messages=list(box);box.close();return messages
def _header(message,name):
 value=str(message.get(name,""))
 try:return str(make_header(decode_header(value)))
 except Exception:return value
def _addresses(message,name):return [address.casefold() for _,address in getaddresses(message.get_all(name,[])) if address]
def _criterion(key,score,evidence,earned=0,total=1):return {"criterion_id":key,"score":round(max(0,min(1,score)),6),"evidence":evidence,"method":"deterministic","raw":{"earned_units":round(earned,6),"total_units":round(total,6)}}
def _fatal(errors,hashes):return {"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":hashes or ["0"*64],"criterion_results":[_criterion(k,0,"；".join(errors),0,0) for k in (*POSITIVE_WEIGHTS,*PENALTY_WEIGHTS)],"reward":0.0,"passed":False,"hurdle_failures":["R001"],"penalties":[]}
def _core(artifact):
 errors=[]
 if not artifact.is_dir():return False,["输出目录不存在"]
 for name in REQUIRED:
  path=artifact/name
  if not path.is_file() or path.stat().st_size==0:errors.append(f"核心交付缺失或为空：{name}");continue
  try:
   if name==XLSX:load_workbook(path,read_only=True,data_only=False).close()
   elif name==DOCX:Document(path)
   elif name==PPTX:
    if not Presentation(path).slides:raise ValueError("PPT无页面")
   elif name in (EMPLOYEE,DIRECTOR):
    messages=_messages(path)
    if not messages or not all(m.get("To") and m.get("Subject") for m in messages):raise ValueError("MBOX无有效草稿")
   elif not isinstance(json.loads(path.read_text()),dict):raise ValueError("JSON顶层不是对象")
  except Exception as exc:errors.append(f"核心交付无法解析：{name}: {exc}")
 return not errors,errors
def _ratio(findings):
 earned=sum(bool(ok) for ok,_ in findings);return earned/len(findings),[label for ok,label in findings if not ok],earned,len(findings)
def _table_findings(seen,wanted,label):
 findings=[(len(seen)==len(wanted),f"{label}行数错误")]
 for ri,wrow in enumerate(wanted):
  observed=seen[ri] if ri<len(seen) else ()
  for ci,value in enumerate(wrow):findings.append((ci<len(observed) and observed[ci]==value,f"{label}第{ri+2}行字段{ci+1}错误"))
 return findings
def _leaf_findings(observed,expected,path):
 if isinstance(expected,dict):
  source=observed if isinstance(observed,dict) else {};findings=[(isinstance(observed,dict),f"{path}应为对象")]
  for key,value in expected.items():findings.extend(_leaf_findings(source.get(key),value,f"{path}.{key}"))
  return findings
 if isinstance(expected,list):
  source=observed if isinstance(observed,list) else [];findings=[(isinstance(observed,list) and len(observed)==len(expected),f"{path}列表长度错误")]
  for index,value in enumerate(expected):findings.extend(_leaf_findings(source[index] if index<len(source) else None,value,f"{path}[{index}]"))
  return findings
 return [(observed==expected,f"{path}应为{expected!r}")]
def _source_rows():
 book=load_workbook(V1.inp()/"Handbook_Acknowledgment_Tracker.xlsx",read_only=True,data_only=True);rows=_rows(book["Acknowledgments"],2,9,7);book.close();return rows
def _source_positive(artifact):
 book=load_workbook(artifact/XLSX,data_only=True);seen=_rows(book["Source_Acknowledgments"],2,9,7);book.close();findings=_table_findings(seen,_source_rows(),"Source_Acknowledgments")
 for filename,digest in V1.H.items():
  path=V1.inp()/filename;findings.append((path.is_file() and _sha256(path)==digest,f"冻结输入{filename}哈希错误"))
 return _ratio(findings)
def _source_policy():
 messages=_messages(V1.inp()/"Handbook_Compliance_Policy.mbox");return [(str(m.get("X-AutomationBench-Message-ID")),str(m.get("From")),_header(m,"Subject"),V1.body_text(m)) for m in messages]
def _trace_positive(artifact):
 rules=[("HB-01","Current acknowledgment date must be >= 2026-03-01"),("HB-02","30-day deadline is 2026-03-31"),("HB-03","45-day escalation date is 2026-04-15"),("HB-04","Reminder employee and CC manager"),("HB-05","At 45 days escalate HR Director"),("HB-06","Approved leave exempt until 7 days after return; Elena until 2026-04-27"),("HB-07","Retain records 3 years")];evidence=[(n,"Frozen 5131 source",h) for n,h in V1.H.items()];book=load_workbook(artifact/XLSX,data_only=True);findings=[]
 for name,wanted,columns in (("Source_Policy",_source_policy(),4),("Policy_Rules",rules,2),("Evidence_Register",evidence,3)):findings.extend(_table_findings(_rows(book[name],2,1+len(wanted),columns),wanted,name))
 book.close();return _ratio(findings)
def _formula_tokens(row,column):
 if column<=7:return (f"Source_Acknowledgments!{'ABCDEFG'[column-1]}{row}",)
 return {8:("DATE(2026,3,1)",),9:("DATE(2026,3,31)",),10:("DATE(2026,4,15)",),11:("AND",f"E{row}",f"H{row}","YES","NO"),12:("SEARCH","On medical leave - returning April 20",f"G{row}","YES","NO"),13:(f"L{row}","DATE(2026,4,27)"),14:(f"L{row}","EXEMPT_LEAVE",f"K{row}",f"E{row}",f"I{row}","COMPLIANT_ON_TIME","COMPLIANT_LATE_ACK","ESCALATE_INVALID_OLD_ACK","ESCALATE_MISSING_ACK"),15:(f"N{row}","2026-04-27","2026-03-31","old handbook does not count","45-day escalation date reached"),16:("LEFT",f"N{row}","ESCALATE","YES","NO"),17:(f"P{row}",f"C{row}"),18:(f"P{row}",f"D{row}"),19:(f"P{row}","DRAFT_NOT_SENT","NOT_REQUIRED"),20:(f"P{row}","YES","NO"),21:(f"P{row}","DRAFT_NOT_SENT","NOT_REQUIRED"),22:('"NO"',),23:("DRAFT_NOT_POSTED",),24:("POLICY_DISTRIBUTION_DATE","30_DAY_DEADLINE","45_DAY_ESCALATION","LEAVE_EXEMPTION"),25:("row=",f"A{row}","employee=",f"B{row}","ack=",f"E{row}","outcome=",f"N{row}")}[column]
def _formula_positive(artifact):
 book=load_workbook(artifact/XLSX,data_only=False);sheet=book["Compliance_Decisions"];formulas=[c for row in sheet.iter_rows(min_row=2,max_row=9,min_col=1,max_col=25) for c in row if isinstance(c.value,str) and c.value.startswith("=")];findings=[(len(formulas)==200,f"Compliance_Decisions公式数应200，实际{len(formulas)}")]
 for row in range(2,10):
  for column in range(1,26):
   cell=sheet.cell(row,column);value=str(cell.value or "");tokens=_formula_tokens(row,column);findings.append((value.startswith("=") and all(t.casefold() in value.casefold() for t in tokens),f"Compliance_Decisions!{cell.coordinate}公式依赖错误"))
 book.close();return _ratio(findings)
def _targets():
 dt=lambda y,m,d:datetime(y,m,d)
 return {
 "Compliant":([(2,"Alice Park",dt(2026,3,15),"Acknowledged","COMPLIANT_ON_TIME","Valid acknowledgment on or before 2026-03-31",None),(5,"Dave Kim",dt(2026,3,25),"Acknowledged","COMPLIANT_ON_TIME","Valid acknowledgment on or before 2026-03-31",None),(8,"Grace Okonkwo",dt(2026,4,2),"Acknowledged","COMPLIANT_LATE_ACK","Valid acknowledgment after deadline; no current reminder",None)],7),
 "Escalations":([(3,"Bob Chen",dt(2026,2,20),"Acknowledged","ESCALATE_INVALID_OLD_ACK","Acknowledgment predates 2026-03-01 distribution; old handbook does not count",None),(4,"Carol Diaz",None,"Not Acknowledged","ESCALATE_MISSING_ACK","No acknowledgment date and 45-day escalation date reached",None),(7,"Femi Adeyemi",dt(2026,2,15),"Acknowledged","ESCALATE_INVALID_OLD_ACK","Acknowledgment predates 2026-03-01 distribution; old handbook does not count",None),(9,"Hugo Fernandez",None,"Not Acknowledged","ESCALATE_MISSING_ACK","No acknowledgment date and 45-day escalation date reached",None)],7),
 "Leave_Exemption":([(6,"Elena Vasquez",None,"Not Acknowledged","EXEMPT_LEAVE","Approved leave; exempt until 2026-04-27",dt(2026,4,27))],7),
 "Employee_Manager_Comms":([(1,"Bob Chen","bob.chen@company.example.com","mgr-b@company.example.com","ESCALATE_INVALID_OLD_ACK","DRAFT_NOT_SENT"),(2,"Carol Diaz","carol.diaz@company.example.com","mgr-a@company.example.com","ESCALATE_MISSING_ACK","DRAFT_NOT_SENT"),(3,"Femi Adeyemi","femi.adeyemi@company.example.com","mgr-b@company.example.com","ESCALATE_INVALID_OLD_ACK","DRAFT_NOT_SENT"),(4,"Hugo Fernandez","hugo.fernandez@company.example.com","mgr-b@company.example.com","ESCALATE_MISSING_ACK","DRAFT_NOT_SENT")],6),
 "HR_Director_Escalation":([("hr-director@company.example.com","Bob Chen; Carol Diaz; Femi Adeyemi; Hugo Fernandez",4,"DRAFT_NOT_SENT")],4),
 "Slack_Control":([("compliance",1,0,0,"8 reviewed; 3 compliant; 4 escalated; 1 leave exempt","DRAFT_NOT_POSTED")],6),
 "Control_Summary":([("Reviewed",8,"Complete"),("Compliant",3,"Alice; Dave; Grace"),("Escalated",4,"Bob; Carol; Femi; Hugo"),("Leave exempt",1,"Elena until 2026-04-27"),("Employee/manager drafts",4,"Not sent"),("HR Director drafts",1,"Not sent"),("Slack drafts",1,"Not posted"),("Real side effects",0,"None")],3)}
def _decision_positive(artifact):
 expected=["Source_Acknowledgments","Source_Policy","Policy_Rules","Evidence_Register","Compliance_Decisions","Compliant","Escalations","Leave_Exemption","Employee_Manager_Comms","HR_Director_Escalation","Slack_Control","Control_Summary"];book=load_workbook(artifact/XLSX,data_only=True);findings=[(book.sheetnames==expected,"Excel十二张工作表名称或顺序错误")]
 for name,(wanted,columns) in _targets().items():findings.extend(_table_findings(_rows(book[name],2,1+len(wanted),columns),wanted,name))
 book.close();return _ratio(findings)
def _leave_positive(artifact):
 book=load_workbook(artifact/XLSX,data_only=True);leave=_rows(book["Leave_Exemption"],2,2,7);escalations=_rows(book["Escalations"],2,5,7);summary=dict((r[0],r[1]) for r in _rows(book["Control_Summary"],2,9,3));book.close();findings=_table_findings(leave,_targets()["Leave_Exemption"][0],"Leave_Exemption");findings.extend([(all(r[1]!="Elena Vasquez" for r in escalations),"Elena不得被升级"),(summary.get("Leave exempt")==1,"休假豁免总数错误")]);return _ratio(findings)
def _employee_positive(artifact):
 messages=_messages(artifact/EMPLOYEE);people=[("Bob Chen","bob.chen@company.example.com","mgr-b@company.example.com","ESCALATE_INVALID_OLD_ACK","Acknowledgment predates 2026-03-01 distribution; old handbook does not count"),("Carol Diaz","carol.diaz@company.example.com","mgr-a@company.example.com","ESCALATE_MISSING_ACK","No acknowledgment date and 45-day escalation date reached"),("Femi Adeyemi","femi.adeyemi@company.example.com","mgr-b@company.example.com","ESCALATE_INVALID_OLD_ACK","Acknowledgment predates 2026-03-01 distribution; old handbook does not count"),("Hugo Fernandez","hugo.fernandez@company.example.com","mgr-b@company.example.com","ESCALATE_MISSING_ACK","No acknowledgment date and 45-day escalation date reached")];findings=[(len(messages)==4,"员工+经理MBOX不是四封逐人草稿")]
 for index,(name,to,cc,outcome,reason) in enumerate(people):
  message=messages[index] if index<len(messages) else None
  if message is None:findings.extend([(False,f"{name}缺少草稿字段") for _ in range(18)]);continue
  body=V1.body_text(message);findings.extend([(_header(message,"From")=="hr-ops@company.example.com",f"{name}发件人错误"),(_addresses(message,"To")==[to],f"{name}To错误"),(_addresses(message,"Cc")==[cc],f"{name}经理Cc错误"),(not _addresses(message,"Bcc"),f"{name}不应有Bcc"),(_header(message,"Subject")=="Required: acknowledge the 2026 employee handbook",f"{name}主题错误"),(_header(message,"X-EvoCowork-State")=="DRAFT_NOT_SENT",f"{name}状态错误"),(_header(message,"X-EvoCowork-Type")=="EMPLOYEE_MANAGER_REMINDER",f"{name}类型错误"),(not message.is_multipart(),f"{name}不应有附件")])
  for token in (name,outcome,reason,"2026-03-01","2026-03-31","2026-04-15","DRAFT_NOT_SENT"):findings.append((token in body,f"{name}正文缺少{token}"))
 return _ratio(findings)
def _director_positive(artifact):
 messages=_messages(artifact/DIRECTOR);findings=[(len(messages)==1,"HR Director MBOX不是一封")];message=messages[0] if messages else None
 if message is None:return _ratio(findings+[(False,"HR Director草稿缺失") for _ in range(20)])
 body=V1.body_text(message);findings.extend([(_header(message,"From")=="hr-ops@company.example.com","HR草稿发件人错误"),(_addresses(message,"To")==["hr-director@company.example.com"],"HR草稿收件人错误"),(_header(message,"Subject")=="Handbook acknowledgment 45-day escalation","HR草稿主题错误"),(_header(message,"X-EvoCowork-State")=="DRAFT_NOT_SENT","HR草稿状态错误"),(_header(message,"X-EvoCowork-Type")=="HR_DIRECTOR_ESCALATION","HR草稿类型错误"),(not _addresses(message,"Cc") and not _addresses(message,"Bcc"),"HR草稿不应Cc/Bcc"),(not message.is_multipart(),"HR草稿不应有附件")])
 for token in ("2026-04-15","Bob Chen","old acknowledgment 2026-02-20","Carol Diaz (missing)","Femi Adeyemi","old acknowledgment 2026-02-15","Hugo Fernandez (missing)","Elena Vasquez is exempt until 2026-04-27","DRAFT_NOT_SENT"):findings.append((token in body,f"HR草稿缺少{token}"))
 return _ratio(findings)
def _expected_slack():return {"channel":"compliance","channel_id":"C_COMP_5131","is_private":False,"state":"DRAFT_NOT_POSTED","draft_count":1,"posted_count":0,"text":"Handbook acknowledgment review: 8 reviewed; 3 compliant; 4 escalated; 1 approved-leave exemption. Follow-up drafts prepared; nothing sent or posted.","names_included":0,"real_side_effects":0}
def _expected_execution():return {"task_id":"automationbench-5131","reviewed":8,"compliant":3,"escalated":4,"leave_exempt":1,"employee_manager_drafts":4,"hr_director_drafts":1,"email_drafts_total":5,"emails_sent":0,"slack_drafts":1,"slack_posts":0,"tracker_updates":0,"real_side_effects":0}
def _controls_positive(artifact):
 findings=[];findings.extend(_leaf_findings(json.loads((artifact/SLACK).read_text()),_expected_slack(),"Handbook_Slack_Summary_Draft"));findings.extend(_leaf_findings(json.loads((artifact/EXECUTION).read_text()),_expected_execution(),"Handbook_Execution_Control"));return _ratio(findings)
def _word_positive(artifact):
 document=Document(artifact/DOCX);rows=[tuple(c.text for c in r.cells) for t in document.tables for r in t.rows];text="\n".join(p.text for p in document.paragraphs)+"\n"+"\n".join(" | ".join(r) for r in rows);findings=[(len(document.sections)==3,"Word不是三分节"),(len(document.tables)==3,"Word不是三张原生表")]
 for row in [("Compliant","3","Alice Park; Dave Kim; Grace Okonkwo"),("Escalate","4","Bob Chen; Carol Diaz; Femi Adeyemi; Hugo Fernandez"),("Leave exempt","1","Elena Vasquez")]:findings.append((row in rows,f"Word总览缺少{row[0]}"))
 for source,outcome in zip(_source_rows(),V1.OUT):
  ack=source[4].strftime("%Y-%m-%d") if source[4] else "blank";reason={"COMPLIANT_ON_TIME":"Valid acknowledgment on or before 2026-03-31","ESCALATE_INVALID_OLD_ACK":"Acknowledgment predates 2026-03-01 distribution; old handbook does not count","ESCALATE_MISSING_ACK":"No acknowledgment date and 45-day escalation date reached","EXEMPT_LEAVE":"Approved leave; exempt until 2026-04-27","COMPLIANT_LATE_ACK":"Valid acknowledgment after deadline; no current reminder"}[outcome];findings.append(((str(source[0]),source[1],ack,source[5],outcome,reason) in rows,f"Word缺少{source[1]}行"))
 for token in ("Bob Chen and Femi Adeyemi have stale Acknowledged labels","Grace Okonkwo acknowledged April 2: late, but valid and complete","No email has been sent, no Slack message posted, and no source record modified"):findings.append((token in text,f"Word缺少{token}"))
 return _ratio(findings)
def _ppt_content(artifact):
 deck=Presentation(artifact/PPTX);texts=[];rows_by=[];tables=[];pictures=0
 for slide in deck.slides:
  parts=[];rows=[];count=0
  for shape in slide.shapes:
   if shape.shape_type==MSO_SHAPE_TYPE.PICTURE:pictures+=1
   if getattr(shape,"has_text_frame",False) and shape.text.strip():parts.append(shape.text.strip())
   if getattr(shape,"has_table",False):count+=1;rows.extend(tuple(c.text for c in r.cells) for r in shape.table.rows)
  texts.append("\n".join(parts));rows_by.append(rows);tables.append(count)
 return texts,rows_by,tables,pictures
def _ppt_positive(artifact):
 texts,rows,tables,pictures=_ppt_content(artifact);findings=[(len(texts)==3,"PPT不是三页"),(tables==[1,1,1],"PPT不是每页一张原生表"),(pictures==0,"PPT使用图片替代关键内容")]
 for row in [("COMPLIANT","3","Alice; Dave; Grace"),("ESCALATE","4","Bob; Carol; Femi; Hugo"),("EXEMPT_LEAVE","1","Elena until 2026-04-27")]:findings.append((row in rows[0],f"PPT第1页缺少{row[0]}"))
 for source,outcome in zip(_source_rows(),V1.OUT):
  ack=source[4].strftime("%Y-%m-%d") if source[4] else "blank";reason={"COMPLIANT_ON_TIME":"Valid acknowledgment on or before 2026-03-31","ESCALATE_INVALID_OLD_ACK":"Acknowledgment predates 2026-03-01 distribution; old handbook does not count","ESCALATE_MISSING_ACK":"No acknowledgment date and 45-day escalation date reached","EXEMPT_LEAVE":"Approved leave; exempt until 2026-04-27","COMPLIANT_LATE_ACK":"Valid acknowledgment after deadline; no current reminder"}[outcome];findings.append(((source[1],ack,source[5],outcome,reason) in rows[1],f"PPT第2页缺少{source[1]}"))
 for row in [("Employee + manager","4","0","DRAFT_NOT_SENT"),("HR Director","1","0","DRAFT_NOT_SENT"),("Slack #compliance","1","0","DRAFT_NOT_POSTED"),("Real side effects","0","0","NONE")]:findings.append((row in rows[2],f"PPT第3页缺少{row[0]}"))
 for index,token in enumerate(("3 compliant · 4 escalated · 1 approved-leave exemption · 0 executed actions","Dates before March 1 refer to the old handbook; late-but-valid and leave cases stay distinct","Four employee+manager reminders, one HR Director escalation, one aggregate Slack draft; nothing executed")):findings.append((token in texts[index],f"PPT第{index+1}页副标题错误"))
 return _ratio(findings)
def _cross_positive(artifact):
 book=load_workbook(artifact/XLSX,data_only=True);comms=_rows(book["Employee_Manager_Comms"],2,5,6);hr=_rows(book["HR_Director_Escalation"],2,2,4)[0];summary=dict((r[0],r[1]) for r in _rows(book["Control_Summary"],2,9,3));book.close();employees=_messages(artifact/EMPLOYEE);director=_messages(artifact/DIRECTOR)[0];slack=json.loads((artifact/SLACK).read_text());execution=json.loads((artifact/EXECUTION).read_text());document=Document(artifact/DOCX);word="\n".join(p.text for p in document.paragraphs)+"\n"+"\n".join(c.text for t in document.tables for r in t.rows for c in r.cells);texts,pptrows,_,_=_ppt_content(artifact);ppt="\n".join(texts)+"\n"+"\n".join(c for sr in pptrows for r in sr for c in r);escalated=["Bob Chen","Carol Diaz","Femi Adeyemi","Hugo Fernandez"]
 findings=[([r[1] for r in comms]==escalated,"Excel升级通信顺序错误"),([_addresses(m,"To")[0] if _addresses(m,"To") else "" for m in employees]==[r[2] for r in comms],"员工MBOX与Excel To冲突"),([_addresses(m,"Cc")[0] if _addresses(m,"Cc") else "" for m in employees]==[r[3] for r in comms],"员工MBOX与Excel经理Cc冲突"),(all(name in V1.body_text(director) for name in escalated) and hr[1]=="; ".join(escalated),"HR升级名单冲突"),(all(summary.get(label)==execution.get(key)==value for label,key,value in [("Reviewed","reviewed",8),("Compliant","compliant",3),("Escalated","escalated",4),("Leave exempt","leave_exempt",1)]),"8/3/4/1跨文件冲突"),(execution.get("employee_manager_drafts")==len(employees)==4 and execution.get("hr_director_drafts")==1 and execution.get("email_drafts_total")==5,"邮件草稿数冲突"),(slack.get("posted_count")==execution.get("slack_posts")==0 and execution.get("emails_sent")==0,"零发送/发布冲突"),(slack.get("real_side_effects")==execution.get("real_side_effects")==0,"零副作用冲突"),(all(name in word and name in ppt for name in V1.NAMES),"Word/PPT未共同覆盖八人")]
 return _ratio(findings)
def _office_positive(artifact):
 book=load_workbook(artifact/XLSX,data_only=False);tables=sum(len(s.tables) for s in book.worksheets);charts=sum(len(s._charts) for s in book.worksheets);conditional=sum(len(i.rules) for i in book["Compliance_Decisions"].conditional_formatting);external=bool(book._external_links);book.close();document=Document(artifact/DOCX);deck=Presentation(artifact/PPTX);ppttables=[sum(bool(getattr(sh,"has_table",False)) for sh in slide.shapes) for slide in deck.slides];pictures=sum(sh.shape_type==MSO_SHAPE_TYPE.PICTURE for slide in deck.slides for sh in slide.shapes);findings=[(tables==12,"Excel不是十二张原生表"),(charts==1,"Excel不是一张原生图表"),(conditional==2,"Excel不是两条条件格式"),(not external,"Excel包含外部链接"),(len(document.tables)==3,"Word不是三张原生表"),(len(deck.slides)==3 and ppttables==[1,1,1] and pictures==0,"PPT三页三表或图片替代错误")]
 for name in (XLSX,DOCX,PPTX):
  safe,detail=V1.package_safe(artifact/name);findings.append((safe,f"{name}不安全：{detail}"))
 for name,pages in ((XLSX,12),(DOCX,3),(PPTX,3)):
  try:count,text=V1.render(artifact/name);findings.extend([(count==pages,f"{name}渲染页数{count}而非{pages}"),(not re.search(r"#(?:VALUE|REF|NAME|DIV/0|N/A)",text,re.I),f"{name}渲染公式错误")])
  except Exception as exc:findings.append((False,f"{name}渲染失败：{type(exc).__name__}: {exc}"))
 return _ratio(findings)
def _decision_severity(artifact):
 book=load_workbook(artifact/XLSX,data_only=True);seen={}
 for sheet in ("Compliant","Escalations","Leave_Exemption"):
  for row in book[sheet].iter_rows(min_row=2,values_only=True):
   if len(row)>4 and row[1]:seen[str(row[1])]=str(row[4])
 book.close();wanted=dict(zip(V1.NAMES,V1.OUT));wrong=[name for name,outcome in wanted.items() if seen.get(name)!=outcome];labels=[f"{name}确认结果错误" for name in wrong];severity=min(1,.35+.16*len(wrong)) if wrong else 0
 if "Elena Vasquez" in wrong:severity=max(severity,.95)
 if any(name in wrong for name in ("Bob Chen","Femi Adeyemi")):severity=max(severity,.85)
 return severity,labels,1.0
def _privacy_severity(artifact):
 employees=_messages(artifact/EMPLOYEE);director=_messages(artifact/DIRECTOR);slack=json.loads((artifact/SLACK).read_text());expected_to=["bob.chen@company.example.com","carol.diaz@company.example.com","femi.adeyemi@company.example.com","hugo.fernandez@company.example.com"];expected_cc=["mgr-b@company.example.com","mgr-a@company.example.com","mgr-b@company.example.com","mgr-b@company.example.com"];severities=[]
 for i,m in enumerate(employees[:4]):
  if _addresses(m,"To")!=[expected_to[i]] or _addresses(m,"Cc")!=[expected_cc[i]]:severities.append((.85,f"第{i+1}封提醒的员工或经理收件人错误"))
 if len(director)!=1 or (director and _addresses(director[0],"To")!=["hr-director@company.example.com"]):severities.append((.9,"45日升级草稿发往错误对象"))
 if slack.get("names_included") not in (0,None) or any(name.casefold() in str(slack.get("text","")).casefold() for name in V1.NAMES):severities.append((.9,"聚合Slack泄露姓名"))
 score=max((s for s,_ in severities),default=0);return score,[x for _,x in severities],1.0
def _execution_severity(artifact):
 employees=_messages(artifact/EMPLOYEE);director=_messages(artifact/DIRECTOR);slack=json.loads((artifact/SLACK).read_text());execution=json.loads((artifact/EXECUTION).read_text());severities=[]
 if any(_header(m,"X-EvoCowork-State")!="DRAFT_NOT_SENT" for m in employees+director) or execution.get("emails_sent") not in (0,None):severities.append((.9,"声称邮件已发送"))
 if slack.get("state") not in ("DRAFT_NOT_POSTED",None) or slack.get("posted_count") not in (0,None) or execution.get("slack_posts") not in (0,None):severities.append((.9,"声称Slack已发布"))
 if execution.get("tracker_updates") not in (0,None):severities.append((.9,"声称源台账已修改"))
 if slack.get("real_side_effects") not in (0,None) or execution.get("real_side_effects") not in (0,None):severities.append((1.0,"声称存在真实副作用"))
 score=max((s for s,_ in severities),default=0);return score,[x for _,x in severities],1.0
def _extra_severity(artifact):
 files={p.name for p in artifact.iterdir() if p.is_file()};dirs=[p.name for p in artifact.iterdir() if p.is_dir()];extras=sorted(files-set(REQUIRED));active=bool(extras or dirs);return (1.0 if active else 0.0),([f"输出含额外文件或子目录：{extras+dirs}"] if active else []),1.0
def _safe_positive(call,artifact):
 try:return call(artifact)
 except Exception as exc:return 0.0,[f"任务专属检查无法完成：{type(exc).__name__}: {exc}"],0,1
def _safe_penalty(call,artifact):
 try:return call(artifact)
 except Exception as exc:return 1.0,[f"严重边界无法核验：{type(exc).__name__}: {exc}"],1.0
def evaluate(artifact):
 ok,errors=_core(artifact);hashes=[_sha256(artifact/n) for n in REQUIRED if (artifact/n).is_file()]
 if not ok:return _fatal(errors,hashes)
 checks={"R002":_source_positive,"R003":_trace_positive,"R004":_formula_positive,"R005":_decision_positive,"R006":_leave_positive,"R007":_employee_positive,"R008":_director_positive,"R009":_controls_positive,"R010":_word_positive,"R011":_ppt_positive,"R012":_cross_positive,"R013":_office_positive};results=[_criterion("R001",1,"七份核心交付存在、非空且可解析",7,7)];scores={"R001":1.0}
 for key in [f"R{i:03d}" for i in range(2,14)]:
  score,failures,earned,total=_safe_positive(checks[key],artifact);results.append(_criterion(key,score,"本题任务专属原子检查通过" if not failures else "未通过："+"；".join(failures),earned,total));scores[key]=score
 for key,call in (("P001",_decision_severity),("P002",_privacy_severity),("P003",_execution_severity),("P004",_extra_severity)):
  score,labels,total=_safe_penalty(call,artifact);results.append(_criterion(key,score,"未触发任务专属严重边界" if not labels else "；".join(labels),score*total,total));scores[key]=score
 total=sum(POSITIVE_WEIGHTS.values());positive=sum(POSITIVE_WEIGHTS[k]*scores[k] for k in POSITIVE_WEIGHTS);penalty=sum(abs(PENALTY_WEIGHTS[k])*scores[k] for k in PENALTY_WEIGHTS);reward=round(max(0,min(1,(positive-penalty)/total)),6);return {"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":hashes,"criterion_results":results,"reward":reward,"passed":reward>=.8,"hurdle_failures":[],"penalties":[{"id":k,"amount":round(abs(PENALTY_WEIGHTS[k])*scores[k]/total,6)} for k in PENALTY_WEIGHTS if scores[k]>0]}
def main():
 p=argparse.ArgumentParser();p.add_argument("--artifact-dir",type=Path,required=True);p.add_argument("--result",type=Path,required=True);p.add_argument("--reward-file",type=Path,required=True);p.add_argument("--scoring-version",choices=("v1","v2"),default="v2");a=p.parse_args()
 try:r=V1.evaluate(a.artifact_dir) if a.scoring_version=="v1" else evaluate(a.artifact_dir)
 except Exception as exc:r=_fatal([f"Judge异常：{type(exc).__name__}: {exc}"],[])
 a.result.parent.mkdir(parents=True,exist_ok=True);a.reward_file.parent.mkdir(parents=True,exist_ok=True);a.result.write_text(json.dumps(r,ensure_ascii=False,indent=2)+"\n");a.reward_file.write_text(f"{float(r['reward']):.6f}\n");print(json.dumps(r,ensure_ascii=False,indent=2))
if __name__=="__main__":main()
