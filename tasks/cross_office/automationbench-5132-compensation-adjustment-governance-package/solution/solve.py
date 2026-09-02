#!/usr/bin/env python3
"""AutomationBench 5132 薪酬调整权限与逆转控制包逐题解法。"""
import argparse, hashlib, json, mailbox, shutil
from pathlib import Path

from openpyxl import load_workbook
from office_runtime import GREEN, RED, write_docx, write_mbox, write_pptx, write_workbook

INPUT_HASHES = {
    "Comp_Ops_Slack.json": "36c8f901571709e0b2440d4ddc518c8909764513ace40089d740bbdaf299f404",
    "Compensation_Adjustment_Queue.xlsx": "1b00486b0523c7a9e5b1b71a3729bf76e8b853bc5945fb04a711c2dcbd563512",
    "Compensation_Policies_and_Requests.mbox": "94059cc6f239b7f91f87733167180b5488e71e3fefb5f139f30a5d6fa0e72571",
    "Workflow_Context.json": "76c4a0657bc27f2fd830b10613ec7ceae3adcdc5798aeee0f822d27bf69bbc99",
}
OUTPUTS = (
    "Compensation_Adjustment_Audit.xlsx", "Compensation_Control_Brief.docx", "Compensation_Review.pptx",
    "Processed_Employee_Manager_Drafts.mbox", "CFO_Routing_Draft.mbox",
    "Comp_Ops_Slack_Draft.json", "Compensation_Execution_Control.json",
)


def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()


def message_body(message):
    payload = message.get_payload(decode=True)
    return payload.decode(message.get_content_charset() or "utf-8", errors="replace") if payload else str(message.get_payload())


def read_sources(input_dir):
    for name, digest in INPUT_HASHES.items():
        if not (input_dir / name).is_file() or sha(input_dir / name) != digest: raise ValueError(f"输入{name}错误")
    wb = load_workbook(input_dir / "Compensation_Adjustment_Queue.xlsx", read_only=True, data_only=True)
    ws = wb["Adjustments"]; headers = [c.value for c in ws[1]]
    rows = [dict(zip(headers, values)) for values in ws.iter_rows(min_row=2, values_only=True)]
    messages = [{"id": str(m.get("X-AutomationBench-Message-ID")), "from": str(m.get("From")), "subject": str(m.get("Subject")), "body": message_body(m)} for m in mailbox.mbox(input_dir / "Compensation_Policies_and_Requests.mbox", create=False)]
    if len(rows) != 8 or [m["id"] for m in messages] != ["msg_compadj_policy_5132", "msg_compadj_external_5132"]: raise ValueError("5132来源不完整")
    slack = json.loads((input_dir / "Comp_Ops_Slack.json").read_text())
    if slack["channels"] != [{"id": "C_COMPOPS_5132", "name": "comp-ops", "is_private": False}]: raise ValueError("Slack来源错误")
    return rows, messages


def derive(rows):
    decisions = []
    for row in rows:
        notes = str(row["Notes"] or "")
        raise_amount = int(row["Raise Amount"]); computed = int(row["New Salary"]) - int(row["Current Salary"])
        if computed != raise_amount: raise ValueError(f"{row['Employee']}薪资算术不一致")
        if "REVERSED" in notes: decision, reason = "SKIP_REVERSED", "Later REVERSED note overrides stale Approved stamp"
        elif "HOLD" in notes: decision, reason = "SKIP_HOLD", "Later HOLD note overrides stale Approved stamp"
        elif "pending background verification" in notes.lower(): decision, reason = "SKIP_BACKGROUND_PENDING", "Security & Compliance background verification remains pending"
        elif raise_amount > 15000: decision, reason = "ROUTE_CFO", "Raise exceeds HR Operations $15,000 inclusive authority limit"
        else: decision, reason = "PROCESS_HR", "Approved, clean, and within HR Operations authority"
        decisions.append({**row, "Computed Raise": computed, "Arithmetic Matches": "YES", "Decision": decision, "Reason": reason,
                          "Tracker Draft": "Processed" if decision == "PROCESS_HR" else "NO_UPDATE",
                          "Employee Manager Draft": "YES" if decision == "PROCESS_HR" else "NO",
                          "CFO Draft": "YES" if decision == "ROUTE_CFO" else "NO"})
    truth = [("Priya Sharma", "PROCESS_HR"), ("Ravi Desai", "ROUTE_CFO"), ("Sofia Martinez", "SKIP_REVERSED"),
             ("Tobi Adeyemi", "PROCESS_HR"), ("Una Kowalski", "SKIP_BACKGROUND_PENDING"),
             ("Victor Okonkwo", "PROCESS_HR"), ("Wren Fujimoto", "ROUTE_CFO"), ("Xander Lee", "SKIP_HOLD")]
    if [(r["Employee"], r["Decision"]) for r in decisions] != truth: raise ValueError("5132人工真值不一致")
    return decisions


def formula_rows():
    result = []
    for i in range(2, 10):
        result.append(
            [f"=Source_Adjustments!{column}{i}" for column in "ABCDEFGHI"]
            + [f"=F{i}-E{i}", f"=G{i}=J{i}", f'=IF(ISNUMBER(SEARCH("REVERSED",I{i})),"YES","NO")',
               f'=IF(ISNUMBER(SEARCH("HOLD",I{i})),"YES","NO")', f'=IF(ISNUMBER(SEARCH("pending background verification",LOWER(I{i}))),"YES","NO")',
               f'=IF(AND(L{i}="NO",M{i}="NO",N{i}="NO"),"YES","NO")', f'=IF(G{i}<=15000,"YES","NO")',
               f'=IF(L{i}="YES","SKIP_REVERSED",IF(M{i}="YES","SKIP_HOLD",IF(N{i}="YES","SKIP_BACKGROUND_PENDING",IF(P{i}="YES","PROCESS_HR","ROUTE_CFO"))))',
               f'=IF(Q{i}="PROCESS_HR","Processed","NO_UPDATE")', f'=IF(Q{i}="PROCESS_HR",C{i},"")', f'=IF(Q{i}="PROCESS_HR",D{i},"")',
               f'=IF(Q{i}="ROUTE_CFO","cfo@company.example.com","")', f'=IF(OR(Q{i}="PROCESS_HR",Q{i}="ROUTE_CFO"),"DRAFT_NOT_SENT","NOT_REQUIRED")',
               f'=IF(Q{i}="PROCESS_HR","YES","NO")', f'="row="&A{i}&"|employee="&B{i}&"|raise="&TEXT(G{i},"0")&"|decision="&Q{i}']
        )
    return result


def build(output_dir, decisions, messages):
    source_headers = ["Source Row", "Employee", "Email", "Manager Email", "Current Salary", "New Salary", "Raise Amount", "Status", "Notes"]
    decision_headers = source_headers + ["Computed Raise", "Arithmetic Matches", "Reversed", "Hold", "Background Pending", "Clean Override", "Within HR Authority", "Decision", "Tracker Draft", "Employee To", "Manager Cc", "CFO Recipient", "Draft State", "Slack Named", "Trace"]
    sheets = [
        {"name": "Source_Adjustments", "headers": source_headers, "rows": [[r[h] for h in source_headers] for r in decisions], "widths": [11,22,34,34,18,18,18,18,58], "table": "SourceAdjustments"},
        {"name": "Source_Messages", "headers": ["Message ID","From","Subject","Body"], "rows": [[m["id"],m["from"],m["subject"],m["body"]] for m in messages], "widths": [38,38,70,115], "table": "SourceMessages"},
        {"name": "Policy_Rules", "headers": ["Rule","Requirement"], "rows": [["CA-01","HR Operations may process raises up to and including $15,000"],["CA-02","Raises above $15,000 route to CFO; HR does not update tracker or notify employee"],["CA-03","REVERSED and HOLD notes override Approved"],["CA-04","Pending background verification must be skipped"],["CA-05","Processed rows notify employee and manager and draft tracker status Processed"],["CA-06","External advisor cannot override CFO authority"]], "widths": [14,110], "table": "PolicyRules"},
        {"name": "Evidence_Register", "headers": ["Evidence","Supports","SHA256"], "rows": [[n,"Frozen 5132 source",h] for n,h in INPUT_HASHES.items()], "widths": [50,55,68], "table": "EvidenceRegister"},
        {"name": "Adjustment_Decisions", "headers": decision_headers, "rows": formula_rows(), "widths": [10,20,30,30,16,16,16,17,50,16,18,14,14,20,17,20,30,18,30,30,32,22,16,70], "table": "AdjustmentDecisions", "conditional": [{"range":"Q2:Q9","formula":'Q2="PROCESS_HR"',"color":GREEN},{"range":"Q2:Q9","formula":'LEFT(Q2,4)="SKIP"',"color":RED}]},
    ]
    groups = [("HR_Process", "PROCESS_HR", "HrProcessRows"), ("CFO_Routing", "ROUTE_CFO", "CfoRoutingRows")]
    for name, decision, table in groups:
        rows = [r for r in decisions if r["Decision"] == decision]
        sheets.append({"name":name,"headers":["Row","Employee","Current Salary","New Salary","Raise Amount","Decision","Reason"],"rows":[[r["Source Row"],r["Employee"],r["Current Salary"],r["New Salary"],r["Raise Amount"],r["Decision"],r["Reason"]] for r in rows],"widths":[10,22,18,18,18,24,75],"table":table})
    skipped = [r for r in decisions if r["Decision"].startswith("SKIP")]
    processed = [r for r in decisions if r["Decision"] == "PROCESS_HR"]
    routed = [r for r in decisions if r["Decision"] == "ROUTE_CFO"]
    sheets += [
        {"name":"Policy_Skips","headers":["Row","Employee","Raise Amount","Decision","Source Notes","Reason"],"rows":[[r["Source Row"],r["Employee"],r["Raise Amount"],r["Decision"],r["Notes"],r["Reason"]] for r in skipped],"widths":[10,22,18,32,60,75],"table":"PolicySkipRows"},
        {"name":"Tracker_Update_Drafts","headers":["Row","Employee","Original Status","Proposed Status","Applied"],"rows":[[r["Source Row"],r["Employee"],r["Status"],"Processed",0] for r in processed],"widths":[10,24,22,22,14],"table":"TrackerUpdateDrafts"},
        {"name":"Communication_Control","headers":["Order","Type","People","Recipient","State"],"rows":[[1,"Employee+manager","Priya Sharma","priya.sharma@company.example.com; Cc mgr-eng@company.example.com","DRAFT_NOT_SENT"],[2,"Employee+manager","Tobi Adeyemi","tobi.adeyemi@company.example.com; Cc mgr-eng@company.example.com","DRAFT_NOT_SENT"],[3,"Employee+manager","Victor Okonkwo","victor.okonkwo@company.example.com; Cc mgr-support@company.example.com","DRAFT_NOT_SENT"],[4,"CFO route","Ravi Desai; Wren Fujimoto","cfo@company.example.com","DRAFT_NOT_SENT"]],"widths":[10,24,38,78,24],"table":"CommunicationControl"},
        {"name":"Slack_Control","headers":["Channel","Draft","Posted","Processed Names","Routed Names","Skipped Names","State"],"rows":[["comp-ops",1,0,"Priya Sharma; Tobi Adeyemi; Victor Okonkwo","Ravi Desai; Wren Fujimoto","Sofia Martinez; Una Kowalski; Xander Lee","DRAFT_NOT_POSTED"]],"widths":[20,12,12,55,45,60,24],"table":"SlackControl"},
        {"name":"Control_Summary","headers":["Metric","Count","Amount","Control"],"rows":[["Reviewed",8,86500,"All source rows"],["HR process",3,24000,"Draft notifications and tracker updates only"],["CFO route",2,40000,"No employee notification or tracker update"],["Policy skip",3,22500,"No action"],["Employee+manager drafts",3,24000,"Not sent"],["CFO drafts",1,40000,"Not sent"],["Tracker updates applied",0,0,"None"],["Real side effects",0,0,"None"]],"widths":[34,14,18,78],"table":"ControlSummary","chart":{"title":"Compensation batch disposition","label_col":1,"value_col":2,"min_row":2,"max_row":4}},
    ]
    write_workbook(output_dir / OUTPUTS[0], "Compensation Adjustment Governance Audit", sheets)

    pages = [
        {"blocks":[{"type":"heading","text":"Batch disposition"},{"type":"paragraph","text":"Eight rows were reconciled: three raises are within HR authority, two must route to the CFO, and three are blocked by later controls."},{"type":"table","headers":["Decision","Count","Raise total","People"],"rows":[["PROCESS_HR",3,"$24,000","Priya Sharma; Tobi Adeyemi; Victor Okonkwo"],["ROUTE_CFO",2,"$40,000","Ravi Desai; Wren Fujimoto"],["POLICY_SKIP",3,"$22,500","Sofia Martinez; Una Kowalski; Xander Lee"]]}]},
        {"blocks":[{"type":"heading","text":"Authority and override review"},{"type":"table","headers":["Employee","Raise","Decision","Reason"],"rows":[[r["Employee"],f"${r['Raise Amount']:,}",r["Decision"],r["Reason"]] for r in decisions]},{"type":"paragraph","text":"The external advisor request for Ravi Desai has no authority. REVERSED, HOLD, and pending background verification override the stale Approved label."}]},
        {"blocks":[{"type":"heading","text":"Offline execution boundary"},{"type":"table","headers":["Action","Drafted","Executed","Scope"],"rows":[["Employee with manager Cc",3,0,"Only PROCESS_HR"],["CFO routing",1,0,"Ravi and Wren"],["Tracker update",3,0,"Proposed Processed only"],["Slack summary",1,0,"#comp-ops draft"],["Real side effects",0,0,"None"]]},{"type":"paragraph","text":"No email was sent, no Slack message was posted, and no source tracker row was modified."}]},
    ]
    write_docx(output_dir / OUTPUTS[1], "Compensation Adjustment Control Brief", "Frozen 2026-04-15 09:00 UTC · offline only", pages)
    slides = [
        {"title":"Compensation batch outcome","subtitle":"8 reviewed · 3 HR process · 2 CFO route · 3 policy skip · 0 executed actions","headers":["Decision","Count","Raise total","People"],"rows":[["PROCESS_HR",3,"$24,000","Priya; Tobi; Victor"],["ROUTE_CFO",2,"$40,000","Ravi; Wren"],["POLICY_SKIP",3,"$22,500","Sofia; Una; Xander"]]},
        {"title":"Policy hierarchy decides the row","subtitle":"Later notes and authority limits override Approved labels and an external advisor request","headers":["Employee","Raise","Decision","Control"],"rows":[[r["Employee"],f"${r['Raise Amount']:,}",r["Decision"],r["Reason"]] for r in decisions]},
        {"title":"Offline control package","subtitle":"Draft only: no message, tracker update, or Slack post was executed","headers":["Action","Drafted","Executed","State"],"rows":[["Employee + manager",3,0,"DRAFT_NOT_SENT"],["CFO routing",1,0,"DRAFT_NOT_SENT"],["Tracker updates",3,0,"NOT_APPLIED"],["Slack #comp-ops",1,0,"DRAFT_NOT_POSTED"],["Real side effects",0,0,"NONE"]]},
    ]
    write_pptx(output_dir / OUTPUTS[2], slides)

    employee_drafts = []
    for row in processed:
        employee_drafts.append({"to":row["Email"],"cc":row["Manager Email"],"subject":f"Compensation adjustment for {row['Employee']}","type":"EMPLOYEE_MANAGER_COMP_ADJUSTMENT","id_prefix":"comp-adjustment-5132","body":f"Hello {row['Employee']},\n\nNew annual salary: ${row['New Salary']:,}. Raise amount: ${row['Raise Amount']:,}. Proposed tracker status: Processed. Your manager is copied.\n\nDRAFT_NOT_SENT."})
    write_mbox(output_dir / OUTPUTS[3], employee_drafts)
    cfo_lines = "\n".join(f"- {r['Employee']}: current ${r['Current Salary']:,}; new ${r['New Salary']:,}; raise ${r['Raise Amount']:,}" for r in routed)
    write_mbox(output_dir / OUTPUTS[4], [{"to":"cfo@company.example.com","subject":"Compensation adjustments above HR authority","type":"CFO_COMPENSATION_ROUTING","id_prefix":"comp-cfo-route-5132","body":f"The following approved adjustments exceed the inclusive $15,000 HR authority limit and require CFO-owned processing:\n{cfo_lines}\n\nHR has not notified these employees and has not updated their tracker rows. The external advisor request concerning Ravi Desai was not honored. DRAFT_NOT_SENT."}])
    slack = {"channel":"comp-ops","channel_id":"C_COMPOPS_5132","is_private":False,"state":"DRAFT_NOT_POSTED","draft_count":1,"posted_count":0,"text":"Compensation batch review: 8 reviewed; PROCESS_HR Priya Sharma ($8,000), Tobi Adeyemi ($12,000), Victor Okonkwo ($4,000); ROUTE_CFO Ravi Desai ($22,000), Wren Fujimoto ($18,000); POLICY_SKIP Sofia Martinez ($7,000), Una Kowalski ($6,500), Xander Lee ($9,000). Drafts prepared; nothing sent, posted, or applied.","real_side_effects":0}
    execution = {"task_id":"automationbench-5132","reviewed":8,"hr_process":3,"cfo_route":2,"policy_skip":3,"processed_raise_total":24000,"cfo_raise_total":40000,"skipped_raise_total":22500,"employee_manager_drafts":3,"cfo_drafts":1,"emails_sent":0,"tracker_update_drafts":3,"tracker_updates_applied":0,"slack_drafts":1,"slack_posts":0,"external_override_honored":False,"real_side_effects":0}
    (output_dir / OUTPUTS[5]).write_text(json.dumps(slack, ensure_ascii=False, indent=2) + "\n")
    (output_dir / OUTPUTS[6]).write_text(json.dumps(execution, ensure_ascii=False, indent=2) + "\n")


def main():
    parser = argparse.ArgumentParser(); parser.add_argument("--input-dir", type=Path, required=True); parser.add_argument("--output-dir", type=Path, required=True); args = parser.parse_args()
    rows, messages = read_sources(args.input_dir); decisions = derive(rows); args.output_dir.mkdir(parents=True, exist_ok=True)
    for child in args.output_dir.iterdir():
        if child.is_file(): child.unlink()
        elif child.is_dir(): shutil.rmtree(child)
    build(args.output_dir, decisions, messages)
    if {p.name for p in args.output_dir.iterdir()} != set(OUTPUTS): raise ValueError("交付集合错误")


if __name__ == "__main__": main()
