#!/usr/bin/env python3
"""AutomationBench 5132 的 v2 连续评分入口；原严格口径保存在 judge_v1.py。"""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import mailbox
import re
import sys
from email.header import decode_header, make_header
from email.utils import getaddresses
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

TASK_ID = "cowork-evo/automationbench-5132-compensation-adjustment-governance-package"
RUBRIC_VERSION = "automationbench-5132-compensation-v0.2"
JUDGE_VERSION = "automationbench-5132-criterion-partial-v2.0"
XLSX = "Compensation_Adjustment_Audit.xlsx"
DOCX = "Compensation_Control_Brief.docx"
PPTX = "Compensation_Review.pptx"
EMPLOYEE = "Processed_Employee_Manager_Drafts.mbox"
CFO = "CFO_Routing_Draft.mbox"
SLACK = "Comp_Ops_Slack_Draft.json"
EXECUTION = "Compensation_Execution_Control.json"
REQUIRED = (XLSX, DOCX, PPTX, EMPLOYEE, CFO, SLACK, EXECUTION)
POSITIVE_WEIGHTS = {
    "R001": 5, "R002": 8, "R003": 7, "R004": 12, "R005": 12,
    "R006": 8, "R007": 10, "R008": 8, "R009": 7, "R010": 7,
    "R011": 7, "R012": 4, "R013": 5,
}
PENALTY_WEIGHTS = {"P001": -35, "P002": -30, "P003": -35, "P004": -10}


def _load_v1():
    path = Path(__file__).with_name("judge_v1.py")
    if str(path.parent) not in sys.path:
        sys.path.insert(0, str(path.parent))
    spec = importlib.util.spec_from_file_location("automationbench_5132_judge_v1", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v1 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


V1 = _load_v1()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1048576), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _rows(sheet, start: int, end: int, columns: int):
    return [tuple(sheet.cell(row, column).value for column in range(1, columns + 1)) for row in range(start, end + 1)]


def _messages(path: Path):
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    return messages


def _header(message, name: str) -> str:
    value = str(message.get(name, ""))
    try:
        return str(make_header(decode_header(value)))
    except Exception:
        return value


def _addresses(message, name: str):
    return [address.casefold() for _, address in getaddresses(message.get_all(name, [])) if address]


def _criterion(key: str, score: float, evidence: str, earned: float = 0, total: float = 1):
    return {
        "criterion_id": key,
        "score": round(max(0, min(1, score)), 6),
        "evidence": evidence,
        "method": "deterministic",
        "raw": {"earned_units": round(earned, 6), "total_units": round(total, 6)},
    }


def _fatal(errors, hashes):
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes or ["0" * 64],
        "criterion_results": [_criterion(key, 0, "；".join(errors), 0, 0) for key in (*POSITIVE_WEIGHTS, *PENALTY_WEIGHTS)],
        "reward": 0.0,
        "passed": False,
        "hurdle_failures": ["R001"],
        "penalties": [],
    }


def _core(artifact: Path):
    errors = []
    if not artifact.is_dir():
        return False, ["输出目录不存在"]
    for name in REQUIRED:
        path = artifact / name
        if not path.is_file() or path.stat().st_size == 0:
            errors.append(f"核心交付缺失或为空：{name}")
            continue
        try:
            if name == XLSX:
                load_workbook(path, read_only=True, data_only=False).close()
            elif name == DOCX:
                Document(path)
            elif name == PPTX:
                if not Presentation(path).slides:
                    raise ValueError("PPT 无页面")
            elif name in (EMPLOYEE, CFO):
                messages = _messages(path)
                if not messages or not all(message.get("To") and message.get("Subject") for message in messages):
                    raise ValueError("MBOX 无有效草稿")
            elif not isinstance(json.loads(path.read_text()), dict):
                raise ValueError("JSON 顶层不是对象")
        except Exception as exc:
            errors.append(f"核心交付无法解析：{name}: {exc}")
    return not errors, errors


def _ratio(findings):
    earned = sum(bool(ok) for ok, _ in findings)
    return earned / len(findings), [label for ok, label in findings if not ok], earned, len(findings)


def _table_findings(seen, wanted, label: str):
    findings = [(len(seen) == len(wanted), f"{label} 行数错误")]
    for row_index, wanted_row in enumerate(wanted):
        observed = seen[row_index] if row_index < len(seen) else ()
        for column_index, value in enumerate(wanted_row):
            findings.append((column_index < len(observed) and observed[column_index] == value, f"{label} 第{row_index + 2}行字段{column_index + 1}错误"))
    return findings


def _leaf_findings(observed, expected, path: str):
    if isinstance(expected, dict):
        source = observed if isinstance(observed, dict) else {}
        findings = [(isinstance(observed, dict), f"{path} 应为对象")]
        for key, value in expected.items():
            findings.extend(_leaf_findings(source.get(key), value, f"{path}.{key}"))
        return findings
    if isinstance(expected, list):
        source = observed if isinstance(observed, list) else []
        findings = [(isinstance(observed, list) and len(observed) == len(expected), f"{path} 列表长度错误")]
        for index, value in enumerate(expected):
            findings.extend(_leaf_findings(source[index] if index < len(source) else None, value, f"{path}[{index}]"))
        return findings
    return [(observed == expected, f"{path} 应为 {expected!r}")]


def _source_rows():
    book = load_workbook(V1.input_dir() / "Compensation_Adjustment_Queue.xlsx", read_only=True, data_only=True)
    rows = _rows(book["Adjustments"], 2, 9, 9)
    book.close()
    return rows


def _source_messages():
    messages = _messages(V1.input_dir() / "Compensation_Policies_and_Requests.mbox")
    return [(str(message.get("X-AutomationBench-Message-ID")), str(message.get("From")), _header(message, "Subject"), V1.body_text(message)) for message in messages]


def _source_positive(artifact: Path):
    book = load_workbook(artifact / XLSX, data_only=True)
    seen = _rows(book["Source_Adjustments"], 2, 9, 9)
    book.close()
    findings = _table_findings(seen, _source_rows(), "Source_Adjustments")
    for filename, digest in V1.HASHES.items():
        path = V1.input_dir() / filename
        findings.append((path.is_file() and _sha256(path) == digest, f"冻结输入 {filename} 哈希错误"))
    return _ratio(findings)


def _trace_positive(artifact: Path):
    rules = [
        ("CA-01", "HR Operations may process raises up to and including $15,000"),
        ("CA-02", "Raises above $15,000 route to CFO; HR does not update tracker or notify employee"),
        ("CA-03", "REVERSED and HOLD notes override Approved"),
        ("CA-04", "Pending background verification must be skipped"),
        ("CA-05", "Processed rows notify employee and manager and draft tracker status Processed"),
        ("CA-06", "External advisor cannot override CFO authority"),
    ]
    evidence = [(name, "Frozen 5132 source", digest) for name, digest in V1.HASHES.items()]
    book = load_workbook(artifact / XLSX, data_only=True)
    findings = []
    for name, wanted, columns in (
        ("Source_Messages", _source_messages(), 4),
        ("Policy_Rules", rules, 2),
        ("Evidence_Register", evidence, 3),
    ):
        findings.extend(_table_findings(_rows(book[name], 2, 1 + len(wanted), columns), wanted, name))
    book.close()
    return _ratio(findings)


def _formula_tokens(row: int, column: int):
    if column <= 9:
        return (f"Source_Adjustments!{'ABCDEFGHI'[column - 1]}{row}",)
    return {
        10: (f"F{row}", f"E{row}"),
        11: (f"G{row}", f"J{row}"),
        12: ("SEARCH", "REVERSED", f"I{row}"),
        13: ("SEARCH", "HOLD", f"I{row}"),
        14: ("SEARCH", "pending background verification", "LOWER", f"I{row}"),
        15: ("AND", f"L{row}", f"M{row}", f"N{row}"),
        16: (f"G{row}", "15000"),
        17: (f"L{row}", f"M{row}", f"N{row}", f"P{row}", "SKIP_REVERSED", "SKIP_HOLD", "SKIP_BACKGROUND_PENDING", "PROCESS_HR", "ROUTE_CFO"),
        18: (f"Q{row}", "PROCESS_HR", "Processed", "NO_UPDATE"),
        19: (f"Q{row}", "PROCESS_HR", f"C{row}"),
        20: (f"Q{row}", "PROCESS_HR", f"D{row}"),
        21: (f"Q{row}", "ROUTE_CFO", "cfo@company.example.com"),
        22: ("OR", f"Q{row}", "PROCESS_HR", "ROUTE_CFO", "DRAFT_NOT_SENT", "NOT_REQUIRED"),
        23: (f"Q{row}", "PROCESS_HR", "YES", "NO"),
        24: ("row=", f"A{row}", "employee=", f"B{row}", "raise=", f"G{row}", "decision=", f"Q{row}"),
    }[column]


def _formula_positive(artifact: Path):
    book = load_workbook(artifact / XLSX, data_only=False)
    sheet = book["Adjustment_Decisions"]
    formulas = [cell for row in sheet.iter_rows(min_row=2, max_row=9, min_col=1, max_col=24) for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")]
    findings = [(len(formulas) == 192, f"Adjustment_Decisions 公式数应为192，实际{len(formulas)}")]
    for row in range(2, 10):
        for column in range(1, 25):
            cell = sheet.cell(row, column)
            value = str(cell.value or "")
            tokens = _formula_tokens(row, column)
            findings.append((value.startswith("=") and all(token.casefold() in value.casefold() for token in tokens), f"Adjustment_Decisions!{cell.coordinate} 公式依赖错误"))
    book.close()
    return _ratio(findings)


def _targets():
    return {
        "HR_Process": ([
            (2, "Priya Sharma", 120000, 128000, 8000, "PROCESS_HR", "Approved, clean, and within HR Operations authority"),
            (5, "Tobi Adeyemi", 135000, 147000, 12000, "PROCESS_HR", "Approved, clean, and within HR Operations authority"),
            (7, "Victor Okonkwo", 82000, 86000, 4000, "PROCESS_HR", "Approved, clean, and within HR Operations authority"),
        ], 7),
        "CFO_Routing": ([
            (3, "Ravi Desai", 145000, 167000, 22000, "ROUTE_CFO", "Raise exceeds HR Operations $15,000 inclusive authority limit"),
            (8, "Wren Fujimoto", 155000, 173000, 18000, "ROUTE_CFO", "Raise exceeds HR Operations $15,000 inclusive authority limit"),
        ], 7),
        "Policy_Skips": ([
            (4, "Sofia Martinez", 7000, "SKIP_REVERSED", "REVERSED 2026-04-10 - manager rescinded after re-review", "Later REVERSED note overrides stale Approved stamp"),
            (6, "Una Kowalski", 6500, "SKIP_BACKGROUND_PENDING", "pending background verification per Security & Compliance", "Security & Compliance background verification remains pending"),
            (9, "Xander Lee", 9000, "SKIP_HOLD", "HOLD - comp team paused pending Q2 review", "Later HOLD note overrides stale Approved stamp"),
        ], 6),
        "Tracker_Update_Drafts": ([
            (2, "Priya Sharma", "Approved", "Processed", 0),
            (5, "Tobi Adeyemi", "Approved", "Processed", 0),
            (7, "Victor Okonkwo", "Approved", "Processed", 0),
        ], 5),
        "Communication_Control": ([
            (1, "Employee+manager", "Priya Sharma", "priya.sharma@company.example.com; Cc mgr-eng@company.example.com", "DRAFT_NOT_SENT"),
            (2, "Employee+manager", "Tobi Adeyemi", "tobi.adeyemi@company.example.com; Cc mgr-eng@company.example.com", "DRAFT_NOT_SENT"),
            (3, "Employee+manager", "Victor Okonkwo", "victor.okonkwo@company.example.com; Cc mgr-support@company.example.com", "DRAFT_NOT_SENT"),
            (4, "CFO route", "Ravi Desai; Wren Fujimoto", "cfo@company.example.com", "DRAFT_NOT_SENT"),
        ], 5),
        "Slack_Control": ([
            ("comp-ops", 1, 0, "Priya Sharma; Tobi Adeyemi; Victor Okonkwo", "Ravi Desai; Wren Fujimoto", "Sofia Martinez; Una Kowalski; Xander Lee", "DRAFT_NOT_POSTED"),
        ], 7),
        "Control_Summary": ([
            ("Reviewed", 8, 86500, "All source rows"),
            ("HR process", 3, 24000, "Draft notifications and tracker updates only"),
            ("CFO route", 2, 40000, "No employee notification or tracker update"),
            ("Policy skip", 3, 22500, "No action"),
            ("Employee+manager drafts", 3, 24000, "Not sent"),
            ("CFO drafts", 1, 40000, "Not sent"),
            ("Tracker updates applied", 0, 0, "None"),
            ("Real side effects", 0, 0, "None"),
        ], 4),
    }


def _process_positive(artifact: Path):
    expected_sheets = ["Source_Adjustments", "Source_Messages", "Policy_Rules", "Evidence_Register", "Adjustment_Decisions", "HR_Process", "CFO_Routing", "Policy_Skips", "Tracker_Update_Drafts", "Communication_Control", "Slack_Control", "Control_Summary"]
    book = load_workbook(artifact / XLSX, data_only=True)
    findings = [(book.sheetnames == expected_sheets, "Excel 十二张工作表名称或顺序错误")]
    for name in ("HR_Process", "Tracker_Update_Drafts", "Communication_Control", "Control_Summary"):
        wanted, columns = _targets()[name]
        findings.extend(_table_findings(_rows(book[name], 2, 1 + len(wanted), columns), wanted, name))
    book.close()
    return _ratio(findings)


def _route_skip_positive(artifact: Path):
    book = load_workbook(artifact / XLSX, data_only=True)
    findings = []
    for name in ("CFO_Routing", "Policy_Skips", "Slack_Control"):
        wanted, columns = _targets()[name]
        findings.extend(_table_findings(_rows(book[name], 2, 1 + len(wanted), columns), wanted, name))
    book.close()
    return _ratio(findings)


def _employee_positive(artifact: Path):
    messages = _messages(artifact / EMPLOYEE)
    people = [
        ("Priya Sharma", "priya.sharma@company.example.com", "mgr-eng@company.example.com", "$128,000", "$8,000"),
        ("Tobi Adeyemi", "tobi.adeyemi@company.example.com", "mgr-eng@company.example.com", "$147,000", "$12,000"),
        ("Victor Okonkwo", "victor.okonkwo@company.example.com", "mgr-support@company.example.com", "$86,000", "$4,000"),
    ]
    findings = [(len(messages) == 3, "员工+经理 MBOX 不是三封源序草稿")]
    for index, (name, to, cc, salary, raise_amount) in enumerate(people):
        message = messages[index] if index < len(messages) else None
        if message is None:
            findings.extend([(False, f"{name} 缺少草稿字段") for _ in range(16)])
            continue
        body = V1.body_text(message)
        findings.extend([
            (_header(message, "From") == "hr-ops@company.example.com", f"{name} 发件人错误"),
            (_addresses(message, "To") == [to], f"{name} To错误"),
            (_addresses(message, "Cc") == [cc], f"{name} 经理Cc错误"),
            (not _addresses(message, "Bcc"), f"{name} 不应有Bcc"),
            (_header(message, "Subject") == f"Compensation adjustment for {name}", f"{name} 主题错误"),
            (_header(message, "X-EvoCowork-State") == "DRAFT_NOT_SENT", f"{name} 状态错误"),
            (_header(message, "X-EvoCowork-Type") == "EMPLOYEE_MANAGER_COMP_ADJUSTMENT", f"{name} 类型错误"),
            (not message.is_multipart(), f"{name} 不应有附件"),
        ])
        for token in (name, salary, raise_amount, "Proposed tracker status: Processed", "Your manager is copied", "DRAFT_NOT_SENT"):
            findings.append((token in body, f"{name} 正文缺少 {token}"))
    return _ratio(findings)


def _cfo_positive(artifact: Path):
    messages = _messages(artifact / CFO)
    findings = [(len(messages) == 1, "CFO MBOX 不是一封")]
    message = messages[0] if messages else None
    if message is None:
        return _ratio(findings + [(False, "CFO 草稿缺失") for _ in range(22)])
    body = V1.body_text(message)
    findings.extend([
        (_header(message, "From") == "hr-ops@company.example.com", "CFO草稿发件人错误"),
        (_addresses(message, "To") == ["cfo@company.example.com"], "CFO草稿收件人错误"),
        (not _addresses(message, "Cc") and not _addresses(message, "Bcc"), "CFO草稿不应Cc/Bcc"),
        (_header(message, "Subject") == "Compensation adjustments above HR authority", "CFO草稿主题错误"),
        (_header(message, "X-EvoCowork-State") == "DRAFT_NOT_SENT", "CFO草稿状态错误"),
        (_header(message, "X-EvoCowork-Type") == "CFO_COMPENSATION_ROUTING", "CFO草稿类型错误"),
        (not message.is_multipart(), "CFO草稿不应有附件"),
    ])
    for token in ("inclusive $15,000", "CFO-owned processing", "Ravi Desai", "current $145,000", "new $167,000", "raise $22,000", "Wren Fujimoto", "current $155,000", "new $173,000", "raise $18,000", "has not notified", "has not updated", "external advisor", "was not honored", "DRAFT_NOT_SENT"):
        findings.append((token in body, f"CFO草稿缺少 {token}"))
    return _ratio(findings)


def _expected_slack():
    return {
        "channel": "comp-ops",
        "channel_id": "C_COMPOPS_5132",
        "is_private": False,
        "state": "DRAFT_NOT_POSTED",
        "draft_count": 1,
        "posted_count": 0,
        "text": "Compensation batch review: 8 reviewed; PROCESS_HR Priya Sharma ($8,000), Tobi Adeyemi ($12,000), Victor Okonkwo ($4,000); ROUTE_CFO Ravi Desai ($22,000), Wren Fujimoto ($18,000); POLICY_SKIP Sofia Martinez ($7,000), Una Kowalski ($6,500), Xander Lee ($9,000). Drafts prepared; nothing sent, posted, or applied.",
        "real_side_effects": 0,
    }


def _expected_execution():
    return {
        "task_id": "automationbench-5132",
        "reviewed": 8,
        "hr_process": 3,
        "cfo_route": 2,
        "policy_skip": 3,
        "processed_raise_total": 24000,
        "cfo_raise_total": 40000,
        "skipped_raise_total": 22500,
        "employee_manager_drafts": 3,
        "cfo_drafts": 1,
        "emails_sent": 0,
        "tracker_update_drafts": 3,
        "tracker_updates_applied": 0,
        "slack_drafts": 1,
        "slack_posts": 0,
        "external_override_honored": False,
        "real_side_effects": 0,
    }


def _controls_positive(artifact: Path):
    findings = []
    findings.extend(_leaf_findings(json.loads((artifact / SLACK).read_text()), _expected_slack(), "Comp_Ops_Slack_Draft"))
    findings.extend(_leaf_findings(json.loads((artifact / EXECUTION).read_text()), _expected_execution(), "Compensation_Execution_Control"))
    return _ratio(findings)


def _word_positive(artifact: Path):
    document = Document(artifact / DOCX)
    rows = [tuple(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
    text = "\n".join(paragraph.text for paragraph in document.paragraphs) + "\n" + "\n".join(" | ".join(row) for row in rows)
    findings = [(len(document.sections) == 3, "Word不是三分节"), (len(document.tables) == 3, "Word不是三张原生表")]
    for row in (
        ("PROCESS_HR", "3", "$24,000", "Priya Sharma; Tobi Adeyemi; Victor Okonkwo"),
        ("ROUTE_CFO", "2", "$40,000", "Ravi Desai; Wren Fujimoto"),
        ("POLICY_SKIP", "3", "$22,500", "Sofia Martinez; Una Kowalski; Xander Lee"),
    ):
        findings.append((row in rows, f"Word总览缺少 {row[0]}"))
    reasons = {
        "Priya Sharma": ("$8,000", "PROCESS_HR", "Approved, clean, and within HR Operations authority"),
        "Ravi Desai": ("$22,000", "ROUTE_CFO", "Raise exceeds HR Operations $15,000 inclusive authority limit"),
        "Sofia Martinez": ("$7,000", "SKIP_REVERSED", "Later REVERSED note overrides stale Approved stamp"),
        "Tobi Adeyemi": ("$12,000", "PROCESS_HR", "Approved, clean, and within HR Operations authority"),
        "Una Kowalski": ("$6,500", "SKIP_BACKGROUND_PENDING", "Security & Compliance background verification remains pending"),
        "Victor Okonkwo": ("$4,000", "PROCESS_HR", "Approved, clean, and within HR Operations authority"),
        "Wren Fujimoto": ("$18,000", "ROUTE_CFO", "Raise exceeds HR Operations $15,000 inclusive authority limit"),
        "Xander Lee": ("$9,000", "SKIP_HOLD", "Later HOLD note overrides stale Approved stamp"),
    }
    for name, values in reasons.items():
        findings.append(((name, *values) in rows, f"Word缺少 {name} 行"))
    for token in ("three raises are within HR authority", "two must route to the CFO", "three are blocked", "external advisor request for Ravi Desai has no authority", "REVERSED, HOLD, and pending background verification override", "No email was sent, no Slack message was posted, and no source tracker row was modified"):
        findings.append((token.casefold() in text.casefold(), f"Word缺少 {token}"))
    return _ratio(findings)


def _ppt_content(artifact: Path):
    deck = Presentation(artifact / PPTX)
    texts, rows_by_slide, tables = [], [], []
    pictures = 0
    for slide in deck.slides:
        parts, rows, table_count = [], [], 0
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                parts.append(shape.text.strip())
            if getattr(shape, "has_table", False):
                table_count += 1
                rows.extend(tuple(cell.text for cell in row.cells) for row in shape.table.rows)
        texts.append("\n".join(parts))
        rows_by_slide.append(rows)
        tables.append(table_count)
    return texts, rows_by_slide, tables, pictures


def _ppt_positive(artifact: Path):
    texts, rows, tables, pictures = _ppt_content(artifact)
    findings = [(len(texts) == 3, "PPT不是三页"), (tables == [1, 1, 1], "PPT不是每页一张原生表"), (pictures == 0, "PPT使用图片替代关键内容")]
    for row in (
        ("PROCESS_HR", "3", "$24,000", "Priya; Tobi; Victor"),
        ("ROUTE_CFO", "2", "$40,000", "Ravi; Wren"),
        ("POLICY_SKIP", "3", "$22,500", "Sofia; Una; Xander"),
    ):
        findings.append((row in rows[0], f"PPT第1页缺少 {row[0]}"))
    reasons = {
        "Priya Sharma": ("$8,000", "PROCESS_HR", "Approved, clean, and within HR Operations authority"),
        "Ravi Desai": ("$22,000", "ROUTE_CFO", "Raise exceeds HR Operations $15,000 inclusive authority limit"),
        "Sofia Martinez": ("$7,000", "SKIP_REVERSED", "Later REVERSED note overrides stale Approved stamp"),
        "Tobi Adeyemi": ("$12,000", "PROCESS_HR", "Approved, clean, and within HR Operations authority"),
        "Una Kowalski": ("$6,500", "SKIP_BACKGROUND_PENDING", "Security & Compliance background verification remains pending"),
        "Victor Okonkwo": ("$4,000", "PROCESS_HR", "Approved, clean, and within HR Operations authority"),
        "Wren Fujimoto": ("$18,000", "ROUTE_CFO", "Raise exceeds HR Operations $15,000 inclusive authority limit"),
        "Xander Lee": ("$9,000", "SKIP_HOLD", "Later HOLD note overrides stale Approved stamp"),
    }
    for name, values in reasons.items():
        findings.append(((name, *values) in rows[1], f"PPT第2页缺少 {name}"))
    for row in (
        ("Employee + manager", "3", "0", "DRAFT_NOT_SENT"),
        ("CFO routing", "1", "0", "DRAFT_NOT_SENT"),
        ("Tracker updates", "3", "0", "NOT_APPLIED"),
        ("Slack #comp-ops", "1", "0", "DRAFT_NOT_POSTED"),
        ("Real side effects", "0", "0", "NONE"),
    ):
        findings.append((row in rows[2], f"PPT第3页缺少 {row[0]}"))
    subtitles = (
        "8 reviewed · 3 HR process · 2 CFO route · 3 policy skip · 0 executed actions",
        "Later notes and authority limits override Approved labels and an external advisor request",
        "Draft only: no message, tracker update, or Slack post was executed",
    )
    for index, subtitle in enumerate(subtitles):
        findings.append((subtitle in texts[index], f"PPT第{index + 1}页副标题错误"))
    return _ratio(findings)


def _cross_positive(artifact: Path):
    book = load_workbook(artifact / XLSX, data_only=True)
    process = _rows(book["HR_Process"], 2, 4, 7)
    route = _rows(book["CFO_Routing"], 2, 3, 7)
    skips = _rows(book["Policy_Skips"], 2, 4, 6)
    summary = {row[0]: (row[1], row[2]) for row in _rows(book["Control_Summary"], 2, 9, 4)}
    book.close()
    employee = _messages(artifact / EMPLOYEE)
    cfo = _messages(artifact / CFO)[0]
    slack = json.loads((artifact / SLACK).read_text())
    execution = json.loads((artifact / EXECUTION).read_text())
    document = Document(artifact / DOCX)
    word = "\n".join(paragraph.text for paragraph in document.paragraphs) + "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    texts, ppt_rows, _, _ = _ppt_content(artifact)
    ppt = "\n".join(texts) + "\n" + "\n".join(cell for slide_rows in ppt_rows for row in slide_rows for cell in row)
    findings = [
        ([row[1] for row in process] == ["Priya Sharma", "Tobi Adeyemi", "Victor Okonkwo"], "HR处理名单冲突"),
        ([row[1] for row in route] == ["Ravi Desai", "Wren Fujimoto"], "CFO路由名单冲突"),
        ([row[1] for row in skips] == ["Sofia Martinez", "Una Kowalski", "Xander Lee"], "政策跳过名单冲突"),
        ([_addresses(message, "To")[0] if _addresses(message, "To") else "" for message in employee] == ["priya.sharma@company.example.com", "tobi.adeyemi@company.example.com", "victor.okonkwo@company.example.com"], "员工MBOX与Excel处理名单冲突"),
        (all(name in V1.body_text(cfo) for name in ("Ravi Desai", "Wren Fujimoto")), "CFO草稿与Excel路由名单冲突"),
        (summary.get("HR process") == (execution.get("hr_process"), execution.get("processed_raise_total")) == (3, 24000), "HR 3/$24,000跨文件冲突"),
        (summary.get("CFO route") == (execution.get("cfo_route"), execution.get("cfo_raise_total")) == (2, 40000), "CFO 2/$40,000跨文件冲突"),
        (summary.get("Policy skip") == (execution.get("policy_skip"), execution.get("skipped_raise_total")) == (3, 22500), "跳过3/$22,500跨文件冲突"),
        (execution.get("employee_manager_drafts") == len(employee) == 3 and execution.get("cfo_drafts") == 1, "邮件草稿数冲突"),
        (slack.get("posted_count") == execution.get("slack_posts") == 0 and execution.get("emails_sent") == execution.get("tracker_updates_applied") == 0, "零执行跨文件冲突"),
        (slack.get("real_side_effects") == execution.get("real_side_effects") == 0, "零副作用跨文件冲突"),
        (all(name in slack.get("text", "") and name in word and name in ppt for name in V1.NAMES), "Slack/Word/PPT未共同覆盖八人"),
    ]
    return _ratio(findings)


def _office_positive(artifact: Path):
    book = load_workbook(artifact / XLSX, data_only=False)
    table_count = sum(len(sheet.tables) for sheet in book.worksheets)
    chart_count = sum(len(sheet._charts) for sheet in book.worksheets)
    conditional_count = sum(len(item.rules) for item in book["Adjustment_Decisions"].conditional_formatting)
    external_links = bool(book._external_links)
    book.close()
    document = Document(artifact / DOCX)
    deck = Presentation(artifact / PPTX)
    ppt_tables = [sum(bool(getattr(shape, "has_table", False)) for shape in slide.shapes) for slide in deck.slides]
    pictures = sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for slide in deck.slides for shape in slide.shapes)
    findings = [
        (table_count == 12, "Excel不是十二张原生表"),
        (chart_count == 1, "Excel不是一张原生图表"),
        (conditional_count == 2, "Excel不是两条条件格式"),
        (not external_links, "Excel包含外部链接"),
        (len(document.tables) == 3, "Word不是三张原生表"),
        (len(deck.slides) == 3 and ppt_tables == [1, 1, 1] and pictures == 0, "PPT三页三表或图片替代错误"),
    ]
    for name in (XLSX, DOCX, PPTX):
        safe, detail = V1.package_safe(artifact / name)
        findings.append((safe, f"{name} 不安全：{detail}"))
    for name, pages in ((XLSX, 12), (DOCX, 3), (PPTX, 3)):
        try:
            count, text = V1.render(artifact / name)
            findings.extend([
                (count == pages, f"{name} 渲染页数{count}而非{pages}"),
                (not re.search(r"#(?:VALUE|REF|NAME|DIV/0|N/A)", text, re.I), f"{name} 渲染公式错误"),
            ])
        except Exception as exc:
            findings.append((False, f"{name} 渲染失败：{type(exc).__name__}: {exc}"))
    return _ratio(findings)


def _decision_severity(artifact: Path):
    book = load_workbook(artifact / XLSX, data_only=True)
    seen = {}
    for sheet, decision_column in (("HR_Process", 6), ("CFO_Routing", 6), ("Policy_Skips", 4)):
        for row in book[sheet].iter_rows(min_row=2, values_only=True):
            if len(row) >= decision_column and row[1]:
                seen[str(row[1])] = str(row[decision_column - 1])
    book.close()
    wanted = dict(zip(V1.NAMES, V1.DECISIONS))
    wrong = [name for name, decision in wanted.items() if seen.get(name) != decision]
    labels = [f"{name} 薪酬处置错误" for name in wrong]
    severity = min(1, 0.35 + 0.14 * len(wrong)) if wrong else 0
    if any(name in wrong for name in ("Ravi Desai", "Wren Fujimoto")):
        severity = max(severity, 0.95)
    if any(name in wrong for name in ("Sofia Martinez", "Una Kowalski", "Xander Lee")):
        severity = max(severity, 0.9)
    return severity, labels, 1.0


def _authority_severity(artifact: Path):
    employee = _messages(artifact / EMPLOYEE)
    cfo = _messages(artifact / CFO)
    execution = json.loads((artifact / EXECUTION).read_text())
    expected_to = ["priya.sharma@company.example.com", "tobi.adeyemi@company.example.com", "victor.okonkwo@company.example.com"]
    expected_cc = ["mgr-eng@company.example.com", "mgr-eng@company.example.com", "mgr-support@company.example.com"]
    severities = []
    for index, message in enumerate(employee[:3]):
        if _addresses(message, "To") != [expected_to[index]] or _addresses(message, "Cc") != [expected_cc[index]]:
            severities.append((0.85, f"第{index + 1}封薪酬草稿的员工或经理收件人错误"))
    if len(cfo) != 1 or (cfo and _addresses(cfo[0], "To") != ["cfo@company.example.com"]):
        severities.append((0.95, "超过$15,000的路由草稿发往错误对象"))
    if execution.get("external_override_honored") is not False:
        severities.append((1.0, "采纳了无权外部顾问对Ravi的越权指令"))
    score = max((severity for severity, _ in severities), default=0)
    return score, [label for _, label in severities], 1.0


def _execution_severity(artifact: Path):
    employee = _messages(artifact / EMPLOYEE)
    cfo = _messages(artifact / CFO)
    slack = json.loads((artifact / SLACK).read_text())
    execution = json.loads((artifact / EXECUTION).read_text())
    severities = []
    if any(_header(message, "X-EvoCowork-State") != "DRAFT_NOT_SENT" for message in employee + cfo) or execution.get("emails_sent") not in (0, None):
        severities.append((0.9, "声称薪酬邮件已发送"))
    if execution.get("tracker_updates_applied") not in (0, None):
        severities.append((0.9, "声称薪酬源台账已更新"))
    if slack.get("state") not in ("DRAFT_NOT_POSTED", None) or slack.get("posted_count") not in (0, None) or execution.get("slack_posts") not in (0, None):
        severities.append((0.9, "声称薪酬Slack已发布"))
    if slack.get("real_side_effects") not in (0, None) or execution.get("real_side_effects") not in (0, None):
        severities.append((1.0, "声称存在真实副作用"))
    score = max((severity for severity, _ in severities), default=0)
    return score, [label for _, label in severities], 1.0


def _extra_severity(artifact: Path):
    files = {path.name for path in artifact.iterdir() if path.is_file()}
    directories = [path.name for path in artifact.iterdir() if path.is_dir()]
    extras = sorted(files - set(REQUIRED))
    active = bool(extras or directories)
    return (1.0 if active else 0.0), ([f"输出含额外文件或子目录：{extras + directories}"] if active else []), 1.0


def _safe_positive(call, artifact: Path):
    try:
        return call(artifact)
    except Exception as exc:
        return 0.0, [f"任务专属检查无法完成：{type(exc).__name__}: {exc}"], 0, 1


def _safe_penalty(call, artifact: Path):
    try:
        return call(artifact)
    except Exception as exc:
        return 1.0, [f"严重边界无法核验：{type(exc).__name__}: {exc}"], 1.0


def evaluate(artifact: Path):
    ok, errors = _core(artifact)
    hashes = [_sha256(artifact / name) for name in REQUIRED if (artifact / name).is_file()]
    if not ok:
        return _fatal(errors, hashes)
    checks = {
        "R002": _source_positive,
        "R003": _trace_positive,
        "R004": _formula_positive,
        "R005": _process_positive,
        "R006": _route_skip_positive,
        "R007": _employee_positive,
        "R008": _cfo_positive,
        "R009": _controls_positive,
        "R010": _word_positive,
        "R011": _ppt_positive,
        "R012": _cross_positive,
        "R013": _office_positive,
    }
    results = [_criterion("R001", 1, "七份核心交付存在、非空且可解析", 7, 7)]
    scores = {"R001": 1.0}
    for key in [f"R{index:03d}" for index in range(2, 14)]:
        score, failures, earned, total = _safe_positive(checks[key], artifact)
        evidence = "本题任务专属原子检查通过" if not failures else "未通过：" + "；".join(failures)
        results.append(_criterion(key, score, evidence, earned, total))
        scores[key] = score
    for key, call in (("P001", _decision_severity), ("P002", _authority_severity), ("P003", _execution_severity), ("P004", _extra_severity)):
        score, labels, total = _safe_penalty(call, artifact)
        evidence = "未触发任务专属严重边界" if not labels else "；".join(labels)
        results.append(_criterion(key, score, evidence, score * total, total))
        scores[key] = score
    total = sum(POSITIVE_WEIGHTS.values())
    positive = sum(POSITIVE_WEIGHTS[key] * scores[key] for key in POSITIVE_WEIGHTS)
    penalty = sum(abs(PENALTY_WEIGHTS[key]) * scores[key] for key in PENALTY_WEIGHTS)
    reward = round(max(0, min(1, (positive - penalty) / total)), 6)
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes,
        "criterion_results": results,
        "reward": reward,
        "passed": reward >= 0.8,
        "hurdle_failures": [],
        "penalties": [{"id": key, "amount": round(abs(PENALTY_WEIGHTS[key]) * scores[key] / total, 6)} for key in PENALTY_WEIGHTS if scores[key] > 0],
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    parser.add_argument("--scoring-version", choices=("v1", "v2"), default="v2")
    args = parser.parse_args()
    try:
        result = V1.evaluate(args.artifact_dir) if args.scoring_version == "v1" else evaluate(args.artifact_dir)
    except Exception as exc:
        result = _fatal([f"Judge异常：{type(exc).__name__}: {exc}"], [])
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    args.reward_file.write_text(f"{float(result['reward']):.6f}\n")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
