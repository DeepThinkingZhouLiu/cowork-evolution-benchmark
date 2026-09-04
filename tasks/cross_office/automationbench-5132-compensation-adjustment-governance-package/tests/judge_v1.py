#!/usr/bin/env python3
"""AutomationBench 5132 薪酬调整权限与逆转控制包任务专属 Judge。"""
import argparse, json, mailbox, re
from pathlib import Path

from openpyxl import load_workbook
from office_checks import body_text, docx_text, package_safe, pptx_text, render, sha256

TASK = "cowork-evo/automationbench-5132-compensation-adjustment-governance-package"
HASHES = {"Comp_Ops_Slack.json":"36c8f901571709e0b2440d4ddc518c8909764513ace40089d740bbdaf299f404","Compensation_Adjustment_Queue.xlsx":"1b00486b0523c7a9e5b1b71a3729bf76e8b853bc5945fb04a711c2dcbd563512","Compensation_Policies_and_Requests.mbox":"94059cc6f239b7f91f87733167180b5488e71e3fefb5f139f30a5d6fa0e72571","Workflow_Context.json":"76c4a0657bc27f2fd830b10613ec7ceae3adcdc5798aeee0f822d27bf69bbc99"}
X="Compensation_Adjustment_Audit.xlsx"; D="Compensation_Control_Brief.docx"; P="Compensation_Review.pptx"; E="Processed_Employee_Manager_Drafts.mbox"; F="CFO_Routing_Draft.mbox"; S="Comp_Ops_Slack_Draft.json"; C="Compensation_Execution_Control.json"; REQUIRED={X,D,P,E,F,S,C}
WEIGHTS={f"R{i:03d}":v for i,v in enumerate([6,8,10,10,10,8,9,8,8,8,7,8],1)}; CAPS={k:(0 if k=="R001" else .25) for k in WEIGHTS if k!="R012"}
NAMES=["Priya Sharma","Ravi Desai","Sofia Martinez","Tobi Adeyemi","Una Kowalski","Victor Okonkwo","Wren Fujimoto","Xander Lee"]
DECISIONS=["PROCESS_HR","ROUTE_CFO","SKIP_REVERSED","PROCESS_HR","SKIP_BACKGROUND_PENDING","PROCESS_HR","ROUTE_CFO","SKIP_HOLD"]


class Checks:
    def __init__(self): self.values={k:{"passed":True,"details":[]} for k in WEIGHTS}
    def require(self,key,condition,detail):
        if not condition:
            self.values[key]["passed"]=False
            if detail not in self.values[key]["details"]: self.values[key]["details"].append(detail)
    def ok(self,key): return self.values[key]["passed"]


def task_root(): return Path(__file__).resolve().parents[1]
def input_dir(): return Path("/app/input_files") if Path("/app/input_files").is_dir() else task_root()/"data/input_files"
def rows(ws): return [tuple(c.value for c in row) for row in ws.iter_rows(min_row=2)]


def check_sources(q):
    for name,digest in HASHES.items(): q.require("R001",(input_dir()/name).is_file() and sha256(input_dir()/name)==digest,f"输入{name}")
    wb=load_workbook(input_dir()/"Compensation_Adjustment_Queue.xlsx",read_only=True,data_only=True); source=rows(wb["Adjustments"])
    q.require("R002",[r[1] for r in source]==NAMES and [r[0] for r in source]==list(range(2,10)),"八行来源")
    q.require("R002",[r[6] for r in source]==[8000,22000,7000,12000,6500,4000,18000,9000],"调薪金额来源")
    q.require("R002",source[2][8].startswith("REVERSED") and "pending background verification" in source[4][8] and source[7][8].startswith("HOLD"),"后续覆盖来源")
    messages=list(mailbox.mbox(input_dir()/"Compensation_Policies_and_Requests.mbox",create=False)); text="\n".join(body_text(m) for m in messages)
    q.require("R002",len(messages)==2 and all(x in text for x in ("up to and including $15,000","above $15,000","REVERSED","HOLD","pending background verification","External Advisor")),"权限与干扰邮件")


def check_workbook(path,q):
    wb=load_workbook(path,data_only=False); expected=["Source_Adjustments","Source_Messages","Policy_Rules","Evidence_Register","Adjustment_Decisions","HR_Process","CFO_Routing","Policy_Skips","Tracker_Update_Drafts","Communication_Control","Slack_Control","Control_Summary"]
    q.require("R003",wb.sheetnames==expected,"12表")
    formulas=[c.value for row in wb["Adjustment_Decisions"].iter_rows(min_row=2,max_row=9) for c in row if isinstance(c.value,str) and c.value.startswith("=")]
    q.require("R004",len(formulas)==192,f"公式{len(formulas)}"); joined="\n".join(formulas)
    for token in ("F2-E2","G2<=15000","REVERSED","HOLD","pending background verification","PROCESS_HR","ROUTE_CFO","SKIP_BACKGROUND_PENDING"):
        q.require("R004",token in joined,f"公式缺{token}")
    process=rows(wb["HR_Process"]); route=rows(wb["CFO_Routing"]); skipped=rows(wb["Policy_Skips"])
    q.require("R005",[(r[1],r[4],r[5]) for r in process]==[("Priya Sharma",8000,"PROCESS_HR"),("Tobi Adeyemi",12000,"PROCESS_HR"),("Victor Okonkwo",4000,"PROCESS_HR")],"HR处理3人")
    q.require("R006",[(r[1],r[4],r[5]) for r in route]==[("Ravi Desai",22000,"ROUTE_CFO"),("Wren Fujimoto",18000,"ROUTE_CFO")],"CFO路由2人")
    q.require("R006",[(r[1],r[2],r[3]) for r in skipped]==[("Sofia Martinez",7000,"SKIP_REVERSED"),("Una Kowalski",6500,"SKIP_BACKGROUND_PENDING"),("Xander Lee",9000,"SKIP_HOLD")],"政策跳过3人")
    q.require("R005",[(r[1],r[3],r[4]) for r in rows(wb["Tracker_Update_Drafts"])]==[("Priya Sharma","Processed",0),("Tobi Adeyemi","Processed",0),("Victor Okonkwo","Processed",0)],"台账草稿")


def check_communications(artifact,q):
    employee=list(mailbox.mbox(artifact/E,create=False)); cfo=list(mailbox.mbox(artifact/F,create=False))
    expected_to=["priya.sharma@company.example.com","tobi.adeyemi@company.example.com","victor.okonkwo@company.example.com"]
    expected_cc=["mgr-eng@company.example.com","mgr-eng@company.example.com","mgr-support@company.example.com"]
    q.require("R007",len(employee)==3 and [str(m.get("To")) for m in employee]==expected_to and [str(m.get("Cc")) for m in employee]==expected_cc,"员工+经理草稿")
    for message,name,new,raise_amount in zip(employee,["Priya Sharma","Tobi Adeyemi","Victor Okonkwo"],["$128,000","$147,000","$86,000"],["$8,000","$12,000","$4,000"]):
        text=body_text(message); q.require("R007",all(token in text for token in (name,new,raise_amount,"DRAFT_NOT_SENT")),f"{name}草稿")
    cfo_text=body_text(cfo[0]) if len(cfo)==1 else ""
    q.require("R008",len(cfo)==1 and str(cfo[0].get("To"))=="cfo@company.example.com" and all(token in cfo_text for token in ("Ravi Desai","$22,000","Wren Fujimoto","$18,000","external advisor","not notified","DRAFT_NOT_SENT")),"CFO路由草稿")
    slack=json.loads((artifact/S).read_text()); q.require("R009",slack.get("channel")=="comp-ops" and slack.get("posted_count")==0 and slack.get("state")=="DRAFT_NOT_POSTED" and all(n in slack.get("text","") for n in NAMES),"Slack草稿")
    control=json.loads((artifact/C).read_text()); expected={"reviewed":8,"hr_process":3,"cfo_route":2,"policy_skip":3,"processed_raise_total":24000,"cfo_raise_total":40000,"skipped_raise_total":22500,"employee_manager_drafts":3,"cfo_drafts":1,"emails_sent":0,"tracker_update_drafts":3,"tracker_updates_applied":0,"slack_drafts":1,"slack_posts":0,"external_override_honored":False,"real_side_effects":0}
    q.require("R009",all(control.get(k)==v for k,v in expected.items()),"执行边界")


def check_office(artifact,q):
    text,sections,tables=docx_text(artifact/D); q.require("R010",sections==3 and tables==3,"Word结构")
    for token in ("three raises are within HR authority","two must route to the CFO","three are blocked","Ravi Desai","Wren Fujimoto","REVERSED","pending background verification","Real side effects | 0 | 0 | None"):
        q.require("R010",token.lower() in text.lower(),f"Word缺{token}")
    q.require("R010","emails were sent" not in text.lower() and "tracker rows were updated" not in text.lower(),"Word虚假执行")
    text,slides,tables,pictures=pptx_text(artifact/P); q.require("R011",slides==3 and tables==[1,1,1] and pictures==0,"PPT结构")
    for token in ("3 HR process","2 CFO route","3 policy skip","Ravi","Wren","Sofia","Una","Xander","DRAFT_NOT_POSTED"):
        q.require("R011",token in text,f"PPT缺{token}")
    for name in (X,D,P):
        safe,detail=package_safe(artifact/name); q.require("R012",safe,f"安全{name}:{detail}")


def finish(q,pages):
    hurdles=[k for k in CAPS if not q.ok(k)]; base=sum(weight for key,weight in WEIGHTS.items() if q.ok(key))/100; cap=min([1]+[CAPS[k] for k in hurdles])
    penalty=(.3 if any(not q.ok(k) for k in ("R002","R003","R004","R005","R006")) else 0)+(.3 if any(not q.ok(k) for k in ("R007","R008","R009")) else 0)+(.25 if any(not q.ok(k) for k in ("R001","R010","R011","R012")) else 0)
    reward=round(max(0,min(base,cap)-penalty),6)
    return {"task_id":TASK,"reward":reward,"passed":reward>=.8 and not hurdles,"criteria":q.values,"hurdle_failures":hurdles,"source_truth":{"reviewed":8,"hr_process":3,"cfo_route":2,"policy_skip":3,"emails_sent":0,"real_side_effects":0},"render_pages":pages}


def evaluate(artifact):
    q=Checks(); check_sources(q); files={p.name for p in artifact.iterdir() if p.is_file()} if artifact.is_dir() else set(); q.require("R001",files==REQUIRED,f"集合{files}")
    if files!=REQUIRED: return finish(q,{})
    for function,keys in ((lambda:check_workbook(artifact/X,q),("R003","R004","R005","R006","R012")),(lambda:check_communications(artifact,q),("R007","R008","R009")),(lambda:check_office(artifact,q),("R010","R011","R012"))):
        try: function()
        except Exception as exc:
            for key in keys: q.require(key,False,str(exc))
    pages={}
    for name,expected in ((X,12),(D,3),(P,3)):
        try:
            count,text=render(artifact/name); pages[name]=count; q.require("R012",count==expected,f"页数{name}:{count}"); q.require("R012",not re.search(r"#(?:VALUE|REF|NAME|DIV/0|N/A)",text,re.I),f"公式错{name}")
        except Exception as exc: q.require("R012",False,str(exc))
    return finish(q,pages)


def main():
    parser=argparse.ArgumentParser(); parser.add_argument("--artifact-dir",type=Path,required=True); parser.add_argument("--result",type=Path,required=True); parser.add_argument("--reward-file",type=Path,required=True); args=parser.parse_args()
    result=evaluate(args.artifact_dir); args.result.parent.mkdir(parents=True,exist_ok=True); args.reward_file.parent.mkdir(parents=True,exist_ok=True); args.result.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n"); args.reward_file.write_text(f"{result['reward']:.6f}\n"); print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=="__main__": main()
