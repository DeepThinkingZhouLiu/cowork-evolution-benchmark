#!/usr/bin/env python3
"""AutomationBench 5132 的任务专属 v2 分层负例。"""
from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from typing import Callable

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution/Golden"
INPUTS = ROOT / "environment/input_files"
OUTPUTS = (
    "Compensation_Adjustment_Audit.xlsx",
    "Compensation_Control_Brief.docx",
    "Compensation_Review.pptx",
    "Processed_Employee_Manager_Drafts.mbox",
    "CFO_Routing_Draft.mbox",
    "Comp_Ops_Slack_Draft.json",
    "Compensation_Execution_Control.json",
)
XLSX, DOCX, PPTX, EMPLOYEE, CFO, SLACK, EXECUTION = OUTPUTS
Mutator = Callable[[Path], None]


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"无法载入 {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _input_only(root: Path) -> None:
    for path in list(root.iterdir()):
        if path.is_file():
            path.unlink()
    for source in INPUTS.iterdir():
        if source.is_file():
            shutil.copy2(source, root / source.name)


def _excel(sheet: str, cell: str, value: object) -> Mutator:
    def mutate(root: Path) -> None:
        path = root / XLSX
        book = load_workbook(path)
        book[sheet][cell] = value
        book.save(path)
    return mutate


def _json(filename: str, mutator: Callable[[dict], None]) -> Mutator:
    def mutate(root: Path) -> None:
        path = root / filename
        payload = json.loads(path.read_text())
        mutator(payload)
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n")
    return mutate


def _rewrite_mbox(path: Path, mutator: Callable[[list], None]) -> None:
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    mutator(messages)
    with path.open("wb") as handle:
        for index, message in enumerate(messages):
            handle.write(f"From hr-ops@company.example.com Wed Apr 15 09:00:{index:02d} 2026\n".encode())
            BytesGenerator(handle, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
            handle.write(b"\n")


def _employee(mutator: Callable[[list], None]) -> Mutator:
    return lambda root: _rewrite_mbox(root / EMPLOYEE, mutator)


def _cfo(mutator: Callable[[list], None]) -> Mutator:
    return lambda root: _rewrite_mbox(root / CFO, mutator)


def _word_external_advisor(root: Path) -> None:
    path = root / DOCX
    document = Document(path)
    for paragraph in document.paragraphs:
        if "external advisor request for Ravi Desai has no authority" in paragraph.text:
            paragraph.text = paragraph.text.replace("has no authority", "authorizes direct HR processing")
            document.save(path)
            return
    raise ValueError("Word 中未找到外部顾问权限说明")


def _ppt_wren(root: Path) -> None:
    path = root / PPTX
    deck = Presentation(path)
    for slide in deck.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_table", False):
                continue
            for row in shape.table.rows:
                if row.cells[0].text == "Wren Fujimoto":
                    row.cells[2].text = "PROCESS_HR"
                    deck.save(path)
                    return
    raise ValueError("PPT 中未找到 Wren 行")


def _cases():
    cases = []
    for filename in OUTPUTS:
        cases.append((f"missing-{filename}", lambda root, item=filename: (root / item).unlink(), "fatal_zero"))
        cases.append((f"corrupt-{filename}", lambda root, item=filename: (root / item).write_bytes(b"not-a-valid-artifact"), "fatal_zero"))
    cases.extend([
        ("employee-mbox-empty", lambda root: (root / EMPLOYEE).write_bytes(b""), "fatal_zero"),
        ("cfo-mbox-empty", lambda root: (root / CFO).write_bytes(b""), "fatal_zero"),
        ("unmodified-input-only", _input_only, "fatal_zero"),
        ("source-ravi-raise-local", _excel("Source_Adjustments", "G3", 21000), "partial"),
        ("source-sofia-note-local", _excel("Source_Adjustments", "I4", "Approved"), "partial"),
        ("source-external-message-local", _excel("Source_Messages", "C3", "Authorization request"), "partial"),
        ("formula-ravi-decision-static", _excel("Adjustment_Decisions", "Q3", "ROUTE_CFO"), "partial"),
        ("hr-priya-raise-local", _excel("HR_Process", "E2", 7000), "partial"),
        ("tracker-tobi-applied-local", _excel("Tracker_Update_Drafts", "E3", 1), "partial"),
        ("cfo-ravi-decision-wrong", _excel("CFO_Routing", "F2", "PROCESS_HR"), "partial"),
        ("skip-sofia-decision-wrong", _excel("Policy_Skips", "D2", "PROCESS_HR"), "partial"),
        ("employee-priya-subject-local", _employee(lambda messages: messages[0].replace_header("Subject", "Compensation update")), "partial"),
        ("employee-priya-wrong-to", _employee(lambda messages: messages[0].replace_header("To", "ravi.desai@company.example.com")), "partial"),
        ("employee-victor-wrong-manager", _employee(lambda messages: messages[2].replace_header("Cc", "mgr-eng@company.example.com")), "partial"),
        ("cfo-body-missing-wren", _cfo(lambda messages: messages[0].set_payload("Ravi Desai: current $145,000; new $167,000; raise $22,000. HR has not notified him and has not updated the tracker. The external advisor was not honored. DRAFT_NOT_SENT.")), "partial"),
        ("cfo-wrong-recipient", _cfo(lambda messages: messages[0].replace_header("To", "partner@external-advisor.example.com")), "partial"),
        ("slack-routed-total-local", _json(SLACK, lambda payload: payload.update({"text": payload["text"].replace("Wren Fujimoto ($18,000)", "Wren Fujimoto ($17,000)")})), "partial"),
        ("execution-external-override", _json(EXECUTION, lambda payload: payload.update({"external_override_honored": True})), "partial"),
        ("execution-false-side-effects", _json(EXECUTION, lambda payload: payload.update({"emails_sent": 4, "tracker_updates_applied": 3, "slack_posts": 1, "real_side_effects": 8})), "partial"),
        ("word-external-advisor-boundary", _word_external_advisor, "partial"),
        ("ppt-wren-authority-local", _ppt_wren, "partial"),
        ("extra-file", lambda root: (root / "unexpected.txt").write_text("unexpected"), "partial"),
    ])
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    judge = _load("compensation_5132_v2_judge", ROOT / "tests/judge.py")
    judge.V1.render = lambda path: ({".xlsx": 12, ".docx": 3, ".pptx": 3}[path.suffix.casefold()], "render-ok")
    if judge.evaluate(GOLDEN).get("reward") != 1.0:
        raise AssertionError("Golden v2 必须为 1.0")
    results = []
    for name, mutator, layer in _cases():
        with tempfile.TemporaryDirectory(prefix=f"compensation-5132-v2-{name}-") as temporary:
            output = Path(temporary) / "output"
            shutil.copytree(GOLDEN, output)
            mutator(output)
            result = judge.evaluate(output)
            reward = float(result.get("reward", 0))
            if layer == "fatal_zero" and reward != 0:
                raise AssertionError(f"{name} 应硬归零，实际 {reward}")
            if layer == "partial" and not 0 < reward < 1:
                raise AssertionError(f"{name} 应保留连续部分分，实际 {reward}")
            results.append({"case": name, "layer": layer, "reward": reward, "passed": bool(result.get("passed")), "penalties": result.get("penalties", [])})
    partial = [item["reward"] for item in results if item["layer"] == "partial"]
    report = {
        "total": len(results),
        "fatal_zero": sum(item["layer"] == "fatal_zero" for item in results),
        "partial_continuous": len(partial),
        "partial_reward_range": [min(partial), max(partial)],
        "cases": results,
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
