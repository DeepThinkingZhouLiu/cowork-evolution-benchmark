#!/usr/bin/env python3
"""对 5132 的缺件、错真值、错收件人、虚假执行和 Office 内容变异做反例回放。"""
import json, mailbox, shutil, subprocess, tempfile
import sys
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]; GOLDEN=ROOT/"solution/Golden"; PYTHON=sys.executable


def score(case):
    result=case.parent/(case.name+".json"); reward=case.parent/(case.name+".txt")
    subprocess.run([PYTHON,str(ROOT/"tests/judge.py"),"--artifact-dir",str(case),"--result",str(result),"--reward-file",str(reward)],check=True,capture_output=True,text=True,timeout=240)
    return float(reward.read_text())


def retarget(path,address):
    box=mailbox.mbox(path,create=False); messages=list(box); box.close(); messages[0].replace_header("To",address)
    with path.open("wb") as handle:
        for index,message in enumerate(messages):
            handle.write(f"From x Wed Apr 15 09:00:{index:02d} 2026\n".encode()); BytesGenerator(handle,policy=SMTP,mangle_from_=True,maxheaderlen=0).flatten(message); handle.write(b"\n")


def main():
    with tempfile.TemporaryDirectory(prefix="comp5132-neg-") as temp:
        root=Path(temp); names=["missing","wrong_counts","external_override","wrong_employee","wrong_cfo","slack_posted","formula_removed","word_false","ppt_omit"]
        cases=[]
        for name in names:
            case=root/name; shutil.copytree(GOLDEN,case); cases.append(case)
        (cases[0]/"Compensation_Review.pptx").unlink()
        payload=json.loads((cases[1]/"Compensation_Execution_Control.json").read_text()); payload["hr_process"]=8; (cases[1]/"Compensation_Execution_Control.json").write_text(json.dumps(payload))
        payload=json.loads((cases[2]/"Compensation_Execution_Control.json").read_text()); payload["external_override_honored"]=True; (cases[2]/"Compensation_Execution_Control.json").write_text(json.dumps(payload))
        retarget(cases[3]/"Processed_Employee_Manager_Drafts.mbox","ravi.desai@company.example.com")
        retarget(cases[4]/"CFO_Routing_Draft.mbox","partner@external-advisor.example.com")
        payload=json.loads((cases[5]/"Comp_Ops_Slack_Draft.json").read_text()); payload["posted_count"]=1; payload["state"]="POSTED"; (cases[5]/"Comp_Ops_Slack_Draft.json").write_text(json.dumps(payload))
        workbook=load_workbook(cases[6]/"Compensation_Adjustment_Audit.xlsx"); workbook["Adjustment_Decisions"]["Q3"]="PROCESS_HR"; workbook.save(cases[6]/"Compensation_Adjustment_Audit.xlsx")
        document=Document(cases[7]/"Compensation_Control_Brief.docx"); document.add_paragraph("All tracker rows were updated and emails were sent."); document.save(cases[7]/"Compensation_Control_Brief.docx")
        presentation=Presentation(cases[8]/"Compensation_Review.pptx")
        for slide in presentation.slides:
            for shape in slide.shapes:
                if getattr(shape,"has_text_frame",False): shape.text=shape.text.replace("Wren","Unknown")
                if getattr(shape,"has_table",False):
                    for row in shape.table.rows:
                        for cell in row.cells: cell.text=cell.text.replace("Wren","Unknown")
        presentation.save(cases[8]/"Compensation_Review.pptx")
        rewards={case.name:score(case) for case in cases}; print(json.dumps(rewards,indent=2)); assert all(value==0 for value in rewards.values()),rewards


if __name__=="__main__": main()
