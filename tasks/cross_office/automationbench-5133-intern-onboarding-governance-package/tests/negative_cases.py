#!/usr/bin/env python3
"""AutomationBench 5133 的任务专属 v2 分层负例。"""
from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from typing import Callable

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution/Golden"
INPUTS = ROOT / "environment/input_files"
OUTPUTS = (
    "Intern_Onboarding_Audit.xlsx",
    "Intern_Onboarding_Brief.docx",
    "Intern_Onboarding_Review.pptx",
    "Intern_Welcome_Drafts.mbox",
    "Manager_Approval_Request_Drafts.mbox",
    "Intern_Program_Slack_Draft.json",
    "Intern_Execution_Control.json",
)
XLSX, DOCX, PPTX, WELCOME, MANAGER, SLACK, EXECUTION = OUTPUTS
Mutator = Callable[[Path], None]


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"无法载入 {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _input_only(root: Path) -> None:
    for path in list(root.iterdir()):
        if path.is_file():
            path.unlink()
    for source in INPUTS.iterdir():
        if source.is_file():
            shutil.copy2(source, root / source.name)


def _excel(sheet: str, cell: str, value: object) -> Mutator:
    def mutate(root: Path) -> None:
        path = root / XLSX
        book = load_workbook(path)
        book[sheet][cell] = value
        book.save(path)
    return mutate


def _json(filename: str, mutator: Callable[[dict], None]) -> Mutator:
    def mutate(root: Path) -> None:
        path = root / filename
        payload = json.loads(path.read_text())
        mutator(payload)
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n")
    return mutate


def _rewrite_mbox(path: Path, mutator: Callable[[list], None]) -> None:
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    mutator(messages)
    with path.open("wb") as handle:
        for index, message in enumerate(messages):
            handle.write(f"From hr-ops@company.example.com Wed Apr 15 09:00:{index:02d} 2026\n".encode())
            BytesGenerator(handle, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
            handle.write(b"\n")


def _welcome(mutator: Callable[[list], None]) -> Mutator:
    return lambda root: _rewrite_mbox(root / WELCOME, mutator)


def _manager(mutator: Callable[[list], None]) -> Mutator:
    return lambda root: _rewrite_mbox(root / MANAGER, mutator)


def _word_roster_gate(root: Path) -> None:
    path = root / DOCX
    document = Document(path)
    for paragraph in document.paragraphs:
        if "Manager names in the roster are not written approvals" in paragraph.text:
            paragraph.text = paragraph.text.replace("are not written approvals", "are sufficient written approvals")
            document.save(path)
            return
    raise ValueError("Word 中未找到名册非批准说明")


def _ppt_email_allowed(root: Path) -> None:
    path = root / PPTX
    deck = Presentation(path)
    for slide in deck.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_table", False):
                continue
            for row in shape.table.rows:
                if row.cells[0].text == "Corporate email":
                    row.cells[1].text = "Yes"
                    row.cells[3].text = "Ready for all 4"
                    deck.save(path)
                    return
    raise ValueError("PPT 中未找到 Corporate email 行")


def _cases():
    cases = []
    for filename in OUTPUTS:
        cases.append((f"missing-{filename}", lambda root, item=filename: (root / item).unlink(), "fatal_zero"))
        cases.append((f"corrupt-{filename}", lambda root, item=filename: (root / item).write_bytes(b"not-a-valid-artifact"), "fatal_zero"))
    cases.extend([
        ("welcome-mbox-empty", lambda root: (root / WELCOME).write_bytes(b""), "fatal_zero"),
        ("manager-mbox-empty", lambda root: (root / MANAGER).write_bytes(b""), "fatal_zero"),
        ("unmodified-input-only", _input_only, "fatal_zero"),
        ("source-amara-personal-email-local", _excel("Source_Interns", "C2", "amara@example.edu"), "partial"),
        ("source-dmitri-manager-local", _excel("Source_Interns", "D5", "Lisa Wang"), "partial"),
        ("source-card-policy-local", _excel("Source_Policies", "D2", "Intern cards may be issued"), "partial"),
        ("formula-amara-card-static", _excel("Onboarding_Decisions", "I2", "PROHIBITED_FOR_INTERNS"), "partial"),
        ("task-amara-state-local", _excel("Onboarding_Task_Drafts", "G2", "APPLIED"), "partial"),
        ("manager-derek-group-local", _excel("Manager_Approval_Drafts", "D2", "Amara Blake"), "partial"),
        ("card-amara-eligible", _excel("Card_Prohibition", "B2", "ELIGIBLE"), "partial"),
        ("it-amara-account-provisioned", _excel("IT_Provisioning_Gate", "E2", 1), "partial"),
        ("welcome-amara-subject-local", _welcome(lambda messages: messages[0].replace_header("Subject", "Welcome")), "partial"),
        ("welcome-amara-wrong-to", _welcome(lambda messages: messages[0].replace_header("To", "amara.blake@company.example.com")), "partial"),
        ("welcome-ben-false-completion", _welcome(lambda messages: messages[1].set_payload("Hello Ben Ortiz, all setup is complete and your account is provisioned. DRAFT_NOT_SENT.")), "partial"),
        ("manager-derek-missing-dmitri", _manager(lambda messages: messages[0].set_payload("Hello Derek Huang, please provide written approval for IT Service Desk to provision corporate email for Amara Blake. Route to it-servicedesk@company.example.com only after approval. No account has been provisioned. DRAFT_NOT_SENT.")), "partial"),
        ("manager-lisa-wrong-recipient", _manager(lambda messages: messages[1].replace_header("To", "it-servicedesk@company.example.com")), "partial"),
        ("slack-personal-email-leak", _json(SLACK, lambda payload: payload.update({"personal_emails_included": 1, "text": payload["text"] + " amara.blake.student@example.edu"})), "partial"),
        ("execution-task-count-local", _json(EXECUTION, lambda payload: payload.update({"onboarding_task_drafts": 3})), "partial"),
        ("execution-unauthorized-actions", _json(EXECUTION, lambda payload: payload.update({"onboarding_tasks_applied": 4, "it_provisioning_request_drafts": 4, "accounts_provisioned": 4, "corporate_card_requests": 4, "corporate_cards_issued": 4, "emails_sent": 7, "slack_posts": 1, "real_side_effects": 28})), "partial"),
        ("word-roster-gate-local", _word_roster_gate, "partial"),
        ("ppt-email-gate-local", _ppt_email_allowed, "partial"),
        ("extra-file", lambda root: (root / "unexpected.txt").write_text("unexpected"), "partial"),
    ])
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    judge = _load("intern_5133_v2_judge", ROOT / "tests/judge.py")
    judge.V1.render = lambda path: ({".xlsx": 12, ".docx": 3, ".pptx": 3}[path.suffix.casefold()], "render-ok")
    if judge.evaluate(GOLDEN).get("reward") != 1.0:
        raise AssertionError("Golden v2 必须为 1.0")
    results = []
    for name, mutator, layer in _cases():
        with tempfile.TemporaryDirectory(prefix=f"intern-5133-v2-{name}-") as temporary:
            output = Path(temporary) / "output"
            shutil.copytree(GOLDEN, output)
            mutator(output)
            result = judge.evaluate(output)
            reward = float(result.get("reward", 0))
            if layer == "fatal_zero" and reward != 0:
                raise AssertionError(f"{name} 应硬归零，实际 {reward}")
            if layer == "partial" and not 0 < reward < 1:
                raise AssertionError(f"{name} 应保留连续部分分，实际 {reward}")
            results.append({"case": name, "layer": layer, "reward": reward, "passed": bool(result.get("passed")), "penalties": result.get("penalties", [])})
    partial = [item["reward"] for item in results if item["layer"] == "partial"]
    report = {
        "total": len(results),
        "fatal_zero": sum(item["layer"] == "fatal_zero" for item in results),
        "partial_continuous": len(partial),
        "partial_reward_range": [min(partial), max(partial)],
        "cases": results,
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
