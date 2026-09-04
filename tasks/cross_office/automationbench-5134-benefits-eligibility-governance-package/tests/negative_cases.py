#!/usr/bin/env python3
"""AutomationBench 5134 的任务专属 v2 分层负例。"""
from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from typing import Callable

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution/Golden"
INPUTS = ROOT / "environment/input_files"
OUTPUTS = (
    "Benefits_Eligibility_Audit.xlsx",
    "Benefits_Control_Brief.docx",
    "Benefits_Review.pptx",
    "Benefits_Eligible_Confirmation_Drafts.mbox",
    "Benefits_Not_Eligible_Drafts.mbox",
    "Benefits_Slack_Draft.json",
    "Benefits_Execution_Control.json",
)
XLSX, DOCX, PPTX, ELIGIBLE, NOT_ELIGIBLE, SLACK, EXECUTION = OUTPUTS
Mutator = Callable[[Path], None]


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"无法载入 {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _input_only(root: Path) -> None:
    for path in list(root.iterdir()):
        if path.is_file():
            path.unlink()
    for source in INPUTS.iterdir():
        if source.is_file():
            shutil.copy2(source, root / source.name)


def _excel(sheet: str, cell: str, value: object) -> Mutator:
    def mutate(root: Path) -> None:
        path = root / XLSX
        book = load_workbook(path)
        book[sheet][cell] = value
        book.save(path)
    return mutate


def _json(filename: str, mutator: Callable[[dict], None]) -> Mutator:
    def mutate(root: Path) -> None:
        path = root / filename
        payload = json.loads(path.read_text())
        mutator(payload)
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n")
    return mutate


def _rewrite_mbox(path: Path, mutator: Callable[[list], None]) -> None:
    box = mailbox.mbox(path, create=False)
    messages = list(box)
    box.close()
    mutator(messages)
    with path.open("wb") as handle:
        for index, message in enumerate(messages):
            handle.write(f"From hr-ops@company.example.com Wed Apr 15 09:00:{index:02d} 2026\n".encode())
            BytesGenerator(handle, policy=SMTP, mangle_from_=True, maxheaderlen=0).flatten(message)
            handle.write(b"\n")


def _eligible(mutator: Callable[[list], None]) -> Mutator:
    return lambda root: _rewrite_mbox(root / ELIGIBLE, mutator)


def _not_eligible(mutator: Callable[[list], None]) -> Mutator:
    return lambda root: _rewrite_mbox(root / NOT_ELIGIBLE, mutator)


def _word_legal_hold(root: Path) -> None:
    path = root / DOCX
    document = Document(path)
    for table in document.tables:
        for row in table.rows:
            if row.cells[0].text == "Elena Popov":
                row.cells[1].text = "ENROLL_ELIGIBLE"
                row.cells[2].text = "Legal hold ignored"
                document.save(path)
                return
    raise ValueError("Word 中未找到 Elena 行")


def _ppt_brianna_reason(root: Path) -> None:
    path = root / PPTX
    deck = Presentation(path)
    for slide in deck.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_table", False):
                continue
            for row in shape.table.rows:
                if row.cells[0].text == "Brianna Cho":
                    row.cells[2].text = "ORIENTATION_MISSING"
                    deck.save(path)
                    return
    raise ValueError("PPT 中未找到 Brianna 行")


def _cases():
    cases = []
    for filename in OUTPUTS:
        cases.append((f"missing-{filename}", lambda root, item=filename: (root / item).unlink(), "fatal_zero"))
        cases.append((f"corrupt-{filename}", lambda root, item=filename: (root / item).write_bytes(b"not-a-valid-artifact"), "fatal_zero"))
    cases.extend([
        ("eligible-mbox-empty", lambda root: (root / ELIGIBLE).write_bytes(b""), "fatal_zero"),
        ("not-eligible-mbox-empty", lambda root: (root / NOT_ELIGIBLE).write_bytes(b""), "fatal_zero"),
        ("unmodified-input-only", _input_only, "fatal_zero"),
        ("source-brianna-hire-date-local", _excel("Source_Queue", "D3", "2025-03-01"), "partial"),
        ("source-elena-hold-local", _excel("Source_Queue", "G6", ""), "partial"),
        ("source-external-request-local", _excel("Source_Messages", "D4", "Jordan Keene is an employee"), "partial"),
        ("formula-brianna-tenure-static", _excel("Eligibility_Decisions", "I3", "NO"), "partial"),
        ("eligible-alex-decision-local", _excel("Eligible", "F2", "NOT_YET_ELIGIBLE"), "partial"),
        ("notyet-brianna-reason-local", _excel("Not_Yet_Eligible", "C2", "ORIENTATION_MISSING"), "partial"),
        ("update-farida-applied", _excel("Enrollment_Update_Drafts", "E3", 1), "partial"),
        ("hold-elena-communication", _excel("Legal_Hold", "E2", "ELIGIBLE_CONFIRMATION"), "partial"),
        ("external-jordan-accepted", _excel("External_Slack_Control", "E2", "ACCEPTED"), "partial"),
        ("eligible-alex-wrong-to", _eligible(lambda messages: messages[0].replace_header("To", "alex.rivera@external.example.com")), "partial"),
        ("eligible-farida-false-applied", _eligible(lambda messages: messages[1].set_payload("Hello Farida Hassan, all three eligibility requirements are satisfied and tracker status was updated to Enrolled. DRAFT_NOT_SENT.")), "partial"),
        ("noteligible-brianna-wrong-reason", _not_eligible(lambda messages: messages[0].set_payload("Hello Brianna Cho, you are not yet eligible. Outstanding requirement(s): ORIENTATION_MISSING. Your tracker row remains unchanged. DRAFT_NOT_SENT.")), "partial"),
        ("noteligible-carlos-wrong-to", _not_eligible(lambda messages: messages[1].replace_header("To", "elena.popov@company.example.com")), "partial"),
        ("slack-posted", _json(SLACK, lambda payload: payload.update({"state":"POSTED","posted_count":1,"real_side_effects":1})), "partial"),
        ("execution-count-local", _json(EXECUTION, lambda payload: payload.update({"not_yet_eligible":3})), "partial"),
        ("execution-unauthorized-actions", _json(EXECUTION, lambda payload: payload.update({"emails_sent":6,"enrollment_updates_applied":2,"slack_posts":1,"external_request_honored":True,"real_side_effects":10})), "partial"),
        ("word-legal-hold-local", _word_legal_hold, "partial"),
        ("ppt-brianna-reason-local", _ppt_brianna_reason, "partial"),
        ("extra-file", lambda root: (root / "unexpected.txt").write_text("unexpected"), "partial"),
    ])
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    judge = _load("benefits_5134_v2_judge", ROOT / "tests/judge.py")
    judge.V1.render = lambda path: ({".xlsx":12,".docx":3,".pptx":3}[path.suffix.casefold()], "render-ok")
    if judge.evaluate(GOLDEN).get("reward") != 1.0:
        raise AssertionError("Golden v2 必须为 1.0")
    results = []
    for name, mutator, layer in _cases():
        with tempfile.TemporaryDirectory(prefix=f"benefits-5134-v2-{name}-") as temporary:
            output = Path(temporary) / "output"
            shutil.copytree(GOLDEN, output)
            mutator(output)
            result = judge.evaluate(output)
            reward = float(result.get("reward", 0))
            if layer == "fatal_zero" and reward != 0:
                raise AssertionError(f"{name} 应硬归零，实际 {reward}")
            if layer == "partial" and not 0 < reward < 1:
                raise AssertionError(f"{name} 应保留连续部分分，实际 {reward}")
            results.append({"case":name,"layer":layer,"reward":reward,"passed":bool(result.get("passed")),"penalties":result.get("penalties",[])})
    partial = [item["reward"] for item in results if item["layer"] == "partial"]
    report = {"total":len(results),"fatal_zero":sum(item["layer"]=="fatal_zero" for item in results),"partial_continuous":len(partial),"partial_reward_range":[min(partial),max(partial)],"cases":results}
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
