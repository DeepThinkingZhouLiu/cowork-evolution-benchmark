#!/usr/bin/env python3
import json,mailbox,shutil,subprocess,tempfile
from email.generator import BytesGenerator
from email.policy import SMTP
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1];GOLDEN=ROOT/"solution/Golden";PYTHON="REMOVED_LOCAL_PATH"
def score(case):
 result=case.parent/(case.name+".json");reward=case.parent/(case.name+".txt");subprocess.run([PYTHON,str(ROOT/"tests/judge.py"),"--artifact-dir",str(case),"--result",str(result),"--reward-file",str(reward)],check=True,capture_output=True,text=True,timeout=240);return float(reward.read_text())
def retarget(path,address):
 box=mailbox.mbox(path,create=False);messages=list(box);box.close();messages[0].replace_header("To",address)
 with path.open("wb") as handle:
  for index,message in enumerate(messages):handle.write(f"From x Wed Apr 15 09:00:{index:02d} 2026\n".encode());BytesGenerator(handle,policy=SMTP,mangle_from_=True,maxheaderlen=0).flatten(message);handle.write(b"\n")
def main():
 with tempfile.TemporaryDirectory(prefix="benefits5134-neg-") as temp:
  root=Path(temp);names=["missing","wrong_counts","elena_notified","external_enrolled","wrong_reason","slack_posted","tracker_applied","formula_removed","word_false","ppt_omit"];cases=[]
  for name in names:case=root/name;shutil.copytree(GOLDEN,case);cases.append(case)
  (cases[0]/"Benefits_Review.pptx").unlink()
  payload=json.loads((cases[1]/"Benefits_Execution_Control.json").read_text());payload["eligible"]=7;(cases[1]/"Benefits_Execution_Control.json").write_text(json.dumps(payload))
  retarget(cases[2]/"Benefits_Not_Eligible_Drafts.mbox","elena.popov@company.example.com")
  retarget(cases[3]/"Benefits_Not_Eligible_Drafts.mbox","jordan.keene@external-broker.example.com")
  retarget(cases[4]/"Benefits_Not_Eligible_Drafts.mbox","brianna.cho@company.example.com");path=cases[4]/"Benefits_Not_Eligible_Drafts.mbox";path.write_bytes(path.read_bytes().replace(b"TENURE_LT_90_DAYS",b"ORIENTATION_MISSING"))
  payload=json.loads((cases[5]/"Benefits_Slack_Draft.json").read_text());payload["posted_count"]=1;payload["state"]="POSTED";(cases[5]/"Benefits_Slack_Draft.json").write_text(json.dumps(payload))
  payload=json.loads((cases[6]/"Benefits_Execution_Control.json").read_text());payload["enrollment_updates_applied"]=2;(cases[6]/"Benefits_Execution_Control.json").write_text(json.dumps(payload))
  workbook=load_workbook(cases[7]/"Benefits_Eligibility_Audit.xlsx");workbook["Eligibility_Decisions"]["N6"]="ENROLL_ELIGIBLE";workbook.save(cases[7]/"Benefits_Eligibility_Audit.xlsx")
  document=Document(cases[8]/"Benefits_Control_Brief.docx");document.add_paragraph("All seven are eligible and emails were sent.");document.save(cases[8]/"Benefits_Control_Brief.docx")
  presentation=Presentation(cases[9]/"Benefits_Review.pptx")
  for slide in presentation.slides:
   for shape in slide.shapes:
    if getattr(shape,"has_text_frame",False):shape.text=shape.text.replace("Elena","Unknown")
    if getattr(shape,"has_table",False):
     for row in shape.table.rows:
      for cell in row.cells:cell.text=cell.text.replace("Elena","Unknown")
  presentation.save(cases[9]/"Benefits_Review.pptx")
  rewards={case.name:score(case) for case in cases};print(json.dumps(rewards,indent=2));assert all(value==0 for value in rewards.values()),rewards
if __name__=="__main__":main()
