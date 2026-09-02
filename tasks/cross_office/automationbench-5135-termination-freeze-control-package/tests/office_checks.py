"""Office Judge 的通用解析与渲染安全检查；不包含任务业务真值。"""

from __future__ import annotations

import hashlib
import re
import subprocess
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""): digest.update(chunk)
    return digest.hexdigest()


def body_text(message) -> str:
    payload = message.get_payload(decode=True)
    if payload is None: return str(message.get_payload())
    return payload.decode(message.get_content_charset() or "utf-8", errors="replace")


def package_safe(path: Path) -> tuple[bool, str]:
    try:
        with zipfile.ZipFile(path) as archive:
            forbidden = [name for name in archive.namelist() if name.lower().endswith(("vbaproject.bin", ".exe", ".dll", ".js", ".vbs"))]
            external = [name for name in archive.namelist() if name.endswith(".rels") and 'TargetMode="External"' in archive.read(name).decode("utf-8", errors="ignore")]
        return (not forbidden and not external), f"危险对象={forbidden}; 外部关系={external}"
    except Exception as exc: return False, str(exc)


def docx_text(path: Path) -> tuple[str, int, int]:
    doc = Document(path); rows = [" | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows]
    return "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(rows), len(doc.sections), len(doc.tables)


def pptx_text(path: Path) -> tuple[str, int, list[int], int]:
    prs = Presentation(path); parts = []; tables = []; pictures = 0
    for slide in prs.slides:
        count = 0
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE: pictures += 1
            if getattr(shape, "has_text_frame", False) and shape.text.strip(): parts.append(shape.text)
            if getattr(shape, "has_table", False): count += 1; parts.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
        tables.append(count)
    return "\n".join(parts), len(prs.slides), tables, pictures


def render(path: Path) -> tuple[int, str]:
    with tempfile.TemporaryDirectory(prefix="office-judge-render-") as temporary:
        root = Path(temporary); profile = root / "profile"; output = root / "out"; profile.mkdir(); output.mkdir()
        result = subprocess.run(["libreoffice", "--headless", f"-env:UserInstallation=file://{profile}", "--convert-to", "pdf", "--outdir", str(output), str(path)], capture_output=True, text=True, timeout=180)
        if result.returncode != 0: raise RuntimeError(result.stderr or result.stdout)
        pdf = output / f"{path.stem}.pdf"; info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True, timeout=60, check=True).stdout
        match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
        if not match: raise RuntimeError("pdfinfo 未返回页数")
        text = subprocess.run(["pdftotext", str(pdf), "-"], capture_output=True, text=True, timeout=60, check=True).stdout
        return int(match.group(1)), text
