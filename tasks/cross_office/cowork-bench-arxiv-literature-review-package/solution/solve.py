#!/usr/bin/env python3
"""基于单题输入生成可回放的 JSON、Excel 和 Word Golden 交付物。"""

from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from arxiv_litreview_processor import build_result, load_json


NAVY = "17324D"
BLUE = "2E75B6"
TEAL = "1F7A8C"
PALE_BLUE = "DDEBF7"
PALE_TEAL = "DDEFEF"
PALE_GOLD = "FFF2CC"
WHITE = "FFFFFF"
GRAY = "64748B"
LIGHT_GRAY = "E2E8F0"
RED = "C00000"


def style_header(ws, row: int, start_col: int, end_col: int) -> None:
    for cells in ws.iter_cols(min_col=start_col, max_col=end_col, min_row=row, max_row=row):
        for cell in cells:
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="thin", color=WHITE))
    ws.row_dimensions[row].height = 30


def add_table(ws, ref: str, name: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def configure_sheet(ws, widths: dict[str, float], landscape: bool = True) -> None:
    ws.freeze_panes = "A2"
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = ws.dimensions
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "Cowork-Bench · Offline reproducible artifact"
    ws.oddFooter.center.size = 8
    ws.oddFooter.center.color = GRAY
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=10, color=NAVY)
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_report_xlsx(path: Path, result: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data_Analysis"
    headers = [
        "Paper_ID",
        "Title",
        "Primary_Category",
        "Published",
        "Benchmark_Area",
        "Growth_Pct",
        "Match_Status",
    ]
    ws.append(headers)
    for row in result["data_analysis"]:
        ws.append([row[key] for key in headers])
    style_header(ws, 1, 1, len(headers))
    configure_sheet(ws, {"A": 16, "B": 46, "C": 19, "D": 15, "E": 24, "F": 14, "G": 16})
    for cell in ws["F"][1:]:
        cell.number_format = '0.00"%"'
    for row in range(2, ws.max_row + 1):
        status = ws.cell(row, 7).value
        ws.cell(row, 7).fill = PatternFill("solid", fgColor=PALE_TEAL if status == "Matched" else PALE_GOLD)
    add_table(ws, f"A1:G{ws.max_row}", "LitReviewData")
    ws.conditional_formatting.add(
        f"F2:F{ws.max_row}",
        ColorScaleRule(start_type="min", start_color="FFF2CC", end_type="max", end_color="63BE7B"),
    )

    metrics = wb.create_sheet("Metrics")
    metrics.append(["Metric", "Value"])
    formulas = [
        ("Total_Papers", "=COUNTA(Data_Analysis!$A$2:$A$5)"),
        ("Excluded_Noise_Papers", result["metrics"]["Excluded_Noise_Papers"]),
        ("Benchmark_Matched_Papers", "=COUNT(Data_Analysis!$F$2:$F$5)"),
        ("Avg_Matched_Growth_Pct", "=ROUND(AVERAGE(Data_Analysis!$F$2:$F$5),2)"),
        ("Top_Growth_Area", "=INDEX(Data_Analysis!$E$2:$E$5,MATCH(MAX(Data_Analysis!$F$2:$F$5),Data_Analysis!$F$2:$F$5,0))"),
        ("Benchmark_Coverage_Pct", "=COUNT(Data_Analysis!$F$2:$F$5)/COUNTA(Data_Analysis!$A$2:$A$5)"),
    ]
    for item in formulas:
        metrics.append(item)
    style_header(metrics, 1, 1, 2)
    configure_sheet(metrics, {"A": 34, "B": 25}, landscape=False)
    metrics["B5"].number_format = '0.00"%"'
    metrics["B7"].number_format = "0.00%"
    add_table(metrics, f"A1:B{metrics.max_row}", "LitReviewMetrics")
    for row in range(2, metrics.max_row + 1):
        metrics.cell(row, 1).font = Font(name="Aptos", size=10, bold=True, color=NAVY)

    recs = wb.create_sheet("Recommendations")
    rec_headers = ["Rank", "Priority", "Action", "Area", "Growth_Pct"]
    recs.append(rec_headers)
    for row in result["recommendations"]:
        recs.append([row[key] for key in rec_headers])
    style_header(recs, 1, 1, len(rec_headers))
    configure_sheet(recs, {"A": 9, "B": 14, "C": 62, "D": 24, "E": 15})
    for cell in recs["E"][1:]:
        cell.number_format = '0.00"%"'
    add_table(recs, f"A1:E{recs.max_row}", "LitReviewRecommendations")

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(path)


def write_tracker_xlsx(path: Path, result: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Arxiv Litreview Tracker"
    headers = ["Paper_ID", "Title", "Primary_Category", "Benchmark_Area", "Growth_Pct", "Review_Status"]
    ws.append(headers)
    for row in result["data_analysis"]:
        review_status = "Ready for review" if row["Match_Status"] == "Matched" else "Needs separate benchmark"
        ws.append(
            [
                row["Paper_ID"],
                row["Title"],
                row["Primary_Category"],
                row["Benchmark_Area"],
                row["Growth_Pct"],
                review_status,
            ]
        )
    style_header(ws, 1, 1, len(headers))
    configure_sheet(ws, {"A": 16, "B": 48, "C": 19, "D": 24, "E": 14, "F": 28})
    for cell in ws["E"][1:]:
        cell.number_format = '0.00"%"'
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 6).fill = PatternFill(
            "solid",
            fgColor=PALE_TEAL if ws.cell(row, 4).value else PALE_GOLD,
        )
    add_table(ws, f"A1:F{ws.max_row}", "ArxivLitreviewTracker")
    wb.save(path)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_doc_table(table, widths: list[float]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for index, cell in enumerate(table.rows[0].cells):
        set_cell_shading(cell, NAVY)
        cell.width = Cm(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(8.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    for row in table.rows[1:]:
        for index, cell in enumerate(row.cells):
            cell.width = Cm(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = RGBColor(23, 50, 77)


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(text)


def write_docx(path: Path, result: dict, benchmark: dict) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    section.orientation = WD_ORIENT.PORTRAIT

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].font.color.rgb = RGBColor(23, 50, 77)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    for style_name, size, color in (("Title", 25, NAVY), ("Heading 1", 16, BLUE), ("Heading 2", 12, TEAL)):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(5)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Аналитический обзор Arxiv")
    subtitle = doc.add_paragraph()
    subtitle_run = subtitle.add_run("Воспроизводимый офлайн-пакет по снимку Cowork-Bench")
    subtitle_run.font.size = Pt(11)
    subtitle_run.font.color.rgb = RGBColor.from_string(GRAY)

    doc.add_heading("Краткое резюме", level=1)
    doc.add_paragraph(
        "В исходном снимке arxiv_papers.json находятся 5 записей. "
        "После фильтра по primary_category остались 4 записи из областей cs.; "
        "1 запись quant-ph исключена как заданный в источнике шум. Прозрачные правила сопоставили 3 из 4 релевантных "
        "записей с внешними темами. Средний показатель роста для сопоставленных тем равен 21,67%, "
        "покрытие — 75,00%, а наибольшее значение имеет тема LLMs (45%)."
    )

    doc.add_heading("Источники и метод", level=1)
    doc.add_paragraph(
        "Анализ использует только два фактических машиночитаемых источника: arxiv_papers.json "
        "(локальный снимок записей) и benchmark_data.json (внешний снимок Academic Trends API). "
        "Файлы analysis_guide.md, Analysis_Guidelines.pdf и Report_Template.xlsx использованы как методические и форматные ориентиры."
    )
    method_steps = [
        "Шаг 1. Сохранить записи, где primary_category начинается с cs.; остальные зафиксировать как исключенные.",
        "Шаг 2. В объединённом поле title + summary искать language model, vision transformer и reinforcement learning.",
        "Шаг 3. Подставить Growth_Pct только из строки benchmark с точно совпавшим названием темы.",
        "Шаг 4. Отсортировать записи по Paper_ID, рассчитать метрики и сверить их между JSON, Excel, Word и Tracker.",
    ]
    for step in method_steps:
        add_bullet(doc, step)

    doc.add_heading("Ключевые выводы", level=1)
    table = doc.add_table(rows=1, cols=5)
    for index, header in enumerate(["Paper_ID", "Название", "Категория", "Тема benchmark", "Рост"]):
        table.rows[0].cells[index].text = header
    for row in result["data_analysis"]:
        cells = table.add_row().cells
        cells[0].text = row["Paper_ID"]
        cells[1].text = row["Title"]
        cells[2].text = row["Primary_Category"]
        cells[3].text = row["Benchmark_Area"] or "нет совпадения"
        cells[4].text = f"{row['Growth_Pct']:.0f}%" if row["Growth_Pct"] is not None else "—"
    style_doc_table(table, [2.4, 6.9, 2.7, 3.3, 1.5])

    doc.add_paragraph(
        "Три внешних значения в исходном снимке: LLMs — 45%, Computer Vision — 12%, RL — 8%. "
        "Запись о federated learning осталась в релевантной выборке, но не получила Growth_Pct, поскольку в benchmark_data.json нет соответствующей темы. "
        "Это не ошибка данных, а явно зафиксированное отсутствие сопоставления."
    )

    doc.add_heading("Рекомендации", level=1)
    rec_table = doc.add_table(rows=1, cols=4)
    for index, header in enumerate(["Ранг", "Приоритет", "Область", "Действие"]):
        rec_table.rows[0].cells[index].text = header
    ru_actions = {
        "LLMs": "Подготовить углублённый синтез по сопоставленной теме; основание — наибольший внешний Growth_Pct 45%.",
        "Computer Vision": "Провести целевой сравнительный обзор; основание — внешний Growth_Pct 12%.",
        "RL": "Вести тему в очереди мониторинга; основание — внешний Growth_Pct 8%.",
    }
    for rec in result["recommendations"]:
        cells = rec_table.add_row().cells
        cells[0].text = str(rec["Rank"])
        cells[1].text = rec["Priority"]
        cells[2].text = rec["Area"]
        cells[3].text = ru_actions[rec["Area"]]
    style_doc_table(rec_table, [1.4, 2.7, 3.7, 9.0])

    doc.add_heading("Ограничения", level=1)
    limitations = [
        "В исходных файлах нет счётчиков цитирования; поэтому отчёт не делает выводов о цитатах.",
        "В исходных файлах нет оценки релевантности на уровне статьи; релевантность определяется только правилом по категории.",
        "Нет внутреннего базового показателя, поэтому gap-анализ из методического PDF нельзя рассчитать без дополнительных данных.",
        "Нет временного ряда и периода наблюдения. Growth_Pct — это дословно перенесённый внешний снимок, а не измерение качества или динамики отдельной статьи.",
        "Записи в arxiv_papers.json взяты из тестового снимка Cowork-Bench; данный отчёт не утверждает, что они дополнительно проверены в публичном Arxiv.",
    ]
    for item in limitations:
        add_bullet(doc, item)

    doc.add_heading("Контроль целостности", level=1)
    doc.add_paragraph(
        "Цепочка контроля: 5 входных записей -> 4 релевантных -> 3 сопоставленных с benchmark. "
        "Те же Paper_ID, Benchmark_Area и Growth_Pct записаны в arxiv_litreview_results.json, Lit_Review_Report.xlsx и Arxiv_Litreview_Tracker.xlsx. "
        "Такая схема позволяет повторить расчёт и обнаружить межфайловые расхождения."
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Cowork-Bench · офлайн-пакет · источники и ограничения зафиксированы")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string(GRAY)

    doc.core_properties.title = "Arxiv Litreview Analysis"
    doc.core_properties.subject = "Offline reproducible Cowork-Bench literature-review package"
    doc.core_properties.author = "Cowork Evolution Benchmark"
    doc.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    papers = load_json(args.input_dir / "arxiv_papers.json")
    benchmark = load_json(args.input_dir / "benchmark_data.json")
    result = build_result(papers, benchmark)

    with (args.output_dir / "arxiv_litreview_results.json").open("w", encoding="utf-8") as handle:
        json.dump(result, handle, ensure_ascii=False, indent=2)
        handle.write("\n")

    processor_source = Path(__file__).with_name("arxiv_litreview_processor.py")
    processor_target = args.output_dir / "arxiv_litreview_processor.py"
    if processor_source.resolve() != processor_target.resolve():
        shutil.copy2(processor_source, processor_target)
    write_report_xlsx(args.output_dir / "Lit_Review_Report.xlsx", result)
    write_tracker_xlsx(args.output_dir / "Arxiv_Litreview_Tracker.xlsx", result)
    write_docx(args.output_dir / "Lit_Review_Analysis.docx", result, benchmark)


if __name__ == "__main__":
    main()
