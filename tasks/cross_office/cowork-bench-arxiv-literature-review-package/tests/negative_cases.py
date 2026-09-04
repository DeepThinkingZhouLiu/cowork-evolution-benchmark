#!/usr/bin/env python3
"""Arxiv 文献综述任务的任务专属 v2 致命/局部反例回归。"""

from __future__ import annotations

import importlib.util
import json
import shutil
import sys
import tempfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from openpyxl.styles import PatternFill


TASK_ROOT = Path(__file__).resolve().parents[1]
GOLDEN = TASK_ROOT / "solution"
REQUIRED = (
    "arxiv_litreview_processor.py",
    "arxiv_litreview_results.json",
    "Lit_Review_Report.xlsx",
    "Lit_Review_Analysis.docx",
    "Arxiv_Litreview_Tracker.xlsx",
)


def _load_judge():
    path = Path(__file__).with_name("judge.py")
    spec = importlib.util.spec_from_file_location("arxiv_litreview_judge_v2_for_negatives", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v2 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


JUDGE = _load_judge()


def _copy_golden(target: Path) -> None:
    target.mkdir(parents=True, exist_ok=True)
    for name in REQUIRED:
        shutil.copy2(GOLDEN / name, target / name)


def _json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _write_json(path: Path, value: dict) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _mutate_json_paper_id(root: Path) -> None:
    path = root / "arxiv_litreview_results.json"
    value = _json(path)
    value["data_analysis"][0]["Paper_ID"] = "2301.99999"
    _write_json(path, value)


def _mutate_json_growth(root: Path) -> None:
    path = root / "arxiv_litreview_results.json"
    value = _json(path)
    value["data_analysis"][1]["Growth_Pct"] = 19
    _write_json(path, value)


def _mutate_json_noise_boundary(root: Path) -> None:
    path = root / "arxiv_litreview_results.json"
    value = _json(path)
    noise = value["excluded_papers"].pop()
    value["data_analysis"].append({
        "Paper_ID": noise["Paper_ID"],
        "Title": noise["Title"],
        "Primary_Category": noise["Primary_Category"],
        "Published": "2023-06-01",
        "Benchmark_Area": None,
        "Growth_Pct": None,
        "Match_Status": "Unmatched",
    })
    _write_json(path, value)


def _mutate_json_metric(root: Path) -> None:
    path = root / "arxiv_litreview_results.json"
    value = _json(path)
    value["metrics"]["Benchmark_Coverage_Pct"] = 50
    _write_json(path, value)


def _mutate_json_recommendation(root: Path) -> None:
    path = root / "arxiv_litreview_results.json"
    value = _json(path)
    value["recommendations"][1]["Priority"] = "High"
    _write_json(path, value)


def _mutate_processor_mapping(root: Path) -> None:
    path = root / "arxiv_litreview_processor.py"
    text = path.read_text(encoding="utf-8")
    path.write_text(text.replace('("vision transformer", "Computer Vision")', '("vision encoder", "Computer Vision")'), encoding="utf-8")


def _mutate_processor_hardcoded(root: Path) -> None:
    payload = _json(root / "arxiv_litreview_results.json")
    script = (
        "#!/usr/bin/env python3\n"
        "import argparse,json\nfrom pathlib import Path\n"
        "p=argparse.ArgumentParser();p.add_argument('--papers');p.add_argument('--benchmark');p.add_argument('--output',type=Path);a=p.parse_args()\n"
        f"payload={payload!r}\n"
        "a.output.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+'\\n',encoding='utf-8')\n"
    )
    (root / "arxiv_litreview_processor.py").write_text(script, encoding="utf-8")


def _mutate_report_cell(root: Path, cell: str, value) -> None:
    path = root / "Lit_Review_Report.xlsx"
    book = load_workbook(path)
    book["Data_Analysis"][cell] = value
    book.save(path)
    book.close()


def _mutate_report_id(root: Path) -> None:
    _mutate_report_cell(root, "A2", "2301.99999")


def _mutate_report_area(root: Path) -> None:
    _mutate_report_cell(root, "E3", "RL")


def _mutate_metric_b2(root: Path) -> None:
    path = root / "Lit_Review_Report.xlsx"
    book = load_workbook(path)
    book["Metrics"]["B2"] = 4
    book.save(path)
    book.close()


def _mutate_metric_b6(root: Path) -> None:
    path = root / "Lit_Review_Report.xlsx"
    book = load_workbook(path)
    book["Metrics"]["B6"] = "LLMs"
    book.save(path)
    book.close()


def _mutate_report_recommendation(root: Path) -> None:
    path = root / "Lit_Review_Report.xlsx"
    book = load_workbook(path)
    book["Recommendations"]["E3"] = 99
    book.save(path)
    book.close()


def _mutate_tracker_id(root: Path) -> None:
    path = root / "Arxiv_Litreview_Tracker.xlsx"
    book = load_workbook(path)
    book["Arxiv Litreview Tracker"]["A2"] = "2301.99999"
    book.save(path)
    book.close()


def _mutate_tracker_status(root: Path) -> None:
    path = root / "Arxiv_Litreview_Tracker.xlsx"
    book = load_workbook(path)
    book["Arxiv Litreview Tracker"]["F5"] = "Ready for review"
    book.save(path)
    book.close()


def _mutate_word_heading(root: Path) -> None:
    path = root / "Lit_Review_Analysis.docx"
    document = Document(path)
    for paragraph in document.paragraphs:
        if paragraph.text.strip().casefold() == "краткое резюме":
            paragraph.text = "Обзор"
            break
    document.save(path)


def _mutate_word_method(root: Path) -> None:
    path = root / "Lit_Review_Analysis.docx"
    document = Document(path)
    for paragraph in document.paragraphs:
        if "primary_category" in paragraph.text:
            paragraph.text = paragraph.text.replace("primary_category", "category")
    document.save(path)


def _mutate_word_unsupported(root: Path) -> None:
    path = root / "Lit_Review_Analysis.docx"
    document = Document(path)
    document.add_paragraph("Citation count: 450.")
    document.save(path)


def _mutate_report_format(root: Path) -> None:
    path = root / "Lit_Review_Report.xlsx"
    book = load_workbook(path)
    book["Data_Analysis"]["A1"].fill = PatternFill(fill_type=None)
    book.save(path)
    book.close()


def _mutate_extra_file(root: Path) -> None:
    (root / "debug-notes.txt").write_text("not requested\n", encoding="utf-8")


PARTIAL_MUTATIONS = {
    "json_one_paper_id_wrong": _mutate_json_paper_id,
    "json_one_growth_wrong": _mutate_json_growth,
    "json_noise_boundary_broken": _mutate_json_noise_boundary,
    "json_one_metric_wrong": _mutate_json_metric,
    "json_one_recommendation_wrong": _mutate_json_recommendation,
    "processor_one_mapping_rule_wrong": _mutate_processor_mapping,
    "processor_hardcoded_against_mutation": _mutate_processor_hardcoded,
    "report_one_paper_id_wrong": _mutate_report_id,
    "report_one_area_wrong": _mutate_report_area,
    "metrics_b2_formula_flattened": _mutate_metric_b2,
    "metrics_b6_formula_flattened": _mutate_metric_b6,
    "report_one_recommendation_growth_wrong": _mutate_report_recommendation,
    "tracker_one_paper_id_wrong": _mutate_tracker_id,
    "tracker_one_status_wrong": _mutate_tracker_status,
    "word_one_required_heading_missing": _mutate_word_heading,
    "word_method_token_missing": _mutate_word_method,
    "word_unsupported_citation_claim": _mutate_word_unsupported,
    "report_one_header_fill_missing": _mutate_report_format,
    "unexpected_debug_file": _mutate_extra_file,
}


def main() -> None:
    records = []
    # 先真实渲染 Golden 一次；下列反例只改原子内容/格式且核心包均由 _core 再解析，
    # 避免为 36 个变体重复启动 LibreOffice，R013 沿用本次已验证的渲染基线。
    render_score, render_failures, _, _ = JUDGE._render(GOLDEN)
    assert render_score == 1, render_failures
    JUDGE._render = lambda _artifact: (1.0, [], 7, 7)
    with tempfile.TemporaryDirectory(prefix="cowork-evo-arxiv-negatives-") as raw:
        root = Path(raw)
        for name in REQUIRED:
            for mode in ("missing", "empty", "corrupt"):
                case = root / f"fatal-{mode}-{name}"
                _copy_golden(case)
                target = case / name
                if mode == "missing":
                    target.unlink()
                elif mode == "empty":
                    target.write_bytes(b"")
                else:
                    target.write_bytes(b"not a valid task artifact")
                score = JUDGE.evaluate(case)["reward"]
                assert score == 0, (mode, name, score)
                records.append({"case": f"{mode}:{name}", "class": "fatal", "reward": score})

        empty = root / "fatal-empty-output-dir"
        empty.mkdir()
        score = JUDGE.evaluate(empty)["reward"]
        assert score == 0, ("empty-output-dir", score)
        records.append({"case": "empty-output-dir", "class": "fatal", "reward": score})

        input_only = root / "fatal-input-only"
        shutil.copytree(TASK_ROOT / "data" / "input_files", input_only)
        score = JUDGE.evaluate(input_only)["reward"]
        assert score == 0, ("input-only", score)
        records.append({"case": "input-only", "class": "fatal", "reward": score})

        for name, mutation in PARTIAL_MUTATIONS.items():
            case = root / f"partial-{name}"
            _copy_golden(case)
            mutation(case)
            score = JUDGE.evaluate(case)["reward"]
            assert 0 < score < 1, (name, score)
            records.append({"case": name, "class": "partial", "reward": score})

    fatal = [row["reward"] for row in records if row["class"] == "fatal"]
    partial = [row["reward"] for row in records if row["class"] == "partial"]
    summary = {
        "fatal_count": len(fatal),
        "fatal_unique_rewards": sorted(set(fatal)),
        "partial_count": len(partial),
        "partial_min": min(partial),
        "partial_max": max(partial),
        "cases": records,
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
