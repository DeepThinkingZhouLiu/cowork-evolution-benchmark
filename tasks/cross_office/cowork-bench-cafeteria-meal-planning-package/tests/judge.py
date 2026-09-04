#!/usr/bin/env python3
"""食堂菜单包的 v2 连续评分入口；原严格口径保存在 judge_v1.py。"""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import re
import sys
import tempfile
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from openpyxl import load_workbook

TASK_ID="cowork-evo/cowork-bench-cafeteria-meal-planning-package"
RUBRIC_VERSION="cowork-bench-cafeteria-meal-planning-package-v0.2"
JUDGE_VERSION="cowork-bench-cafeteria-meal-planning-criterion-partial-v2.0"
PROCESSOR="meal_cost_optimizer.py";JSON_NAME="meal_optimization_results.json";XLSX="Weekly_Meal_Budget.xlsx";DOCX="Cafeteria_Meal_Proposal.docx";ICS="meal_prep_calendar.ics"
REQUIRED=(PROCESSOR,JSON_NAME,XLSX,DOCX,ICS)
POSITIVE_WEIGHTS={"R001":5,"R002":7,"R003":7,"R004":6,"R005":8,"R006":7,"R007":8,"R008":7,"R009":6,"R010":7,"R011":6,"R012":4,"R013":3,"R014":4}
PENALTY_WEIGHTS={"R015":-15,"R016":-14,"R017":-14,"R018":-18}
PRICE_FIELDS=["Item","Unit","FreshMart_Price","ValueGrocer_Price","BulkBarn_Price","Cheapest_Store","Cheapest_Price"]
RECIPE_FIELDS=["Recipe_ID","Recipe","Matched_Ingredient_Count","Unpriced_Ingredient_Count","Total_Cost","Cost_Per_Serving"]
PLAN_FIELDS=["Date","Day","Meal","Recipe_ID","Recipe","Estimated_Cost"]


def _load_v1():
    path=Path(__file__).with_name("judge_v1.py")
    if str(path.parent) not in sys.path:sys.path.insert(0,str(path.parent))
    spec=importlib.util.spec_from_file_location("cafeteria_judge_v1",path)
    if spec is None or spec.loader is None:raise RuntimeError("无法载入 v1 Judge")
    module=importlib.util.module_from_spec(spec);sys.modules[spec.name]=module;spec.loader.exec_module(module);return module


V1=_load_v1()


def _sha(path:Path)->str:
    digest=hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda:handle.read(1048576),b""):digest.update(chunk)
    return digest.hexdigest()


def _criterion(key:str,score:float,evidence:str,earned:float=0,total:float=1)->dict[str,Any]:
    return {"criterion_id":key,"score":round(max(0,min(1,score)),6),"evidence":evidence,"method":"deterministic","raw":{"earned_units":round(earned,6),"total_units":round(total,6)}}


def _ratio(findings:Iterable[tuple[bool,str]]):
    rows=list(findings);earned=sum(bool(ok) for ok,_ in rows);return earned/max(1,len(rows)),[label for ok,label in rows if not ok],earned,len(rows)


def _same(actual:Any,expected:Any)->bool:
    if isinstance(expected,(int,float)) and not isinstance(expected,bool):return V1.close(actual,expected)
    if expected is None:return actual in (None,"")
    return str(actual).strip()==str(expected).strip()


def _rows(actual:Any,expected:list[dict],fields:list[str],label:str):
    seen=actual if isinstance(actual,list) else [];findings=[(isinstance(actual,list),f"{label} 应为数组"),(len(seen)==len(expected),f"{label} 行数错误")]
    for index,wanted in enumerate(expected):
        row=seen[index] if index<len(seen) and isinstance(seen[index],dict) else {}
        for field in fields:findings.append((_same(row.get(field),wanted.get(field)),f"{label} 第{index+1}行 {field} 错误"))
    return findings


def _recipe_findings(actual:Any,expected:list[dict],label:str):
    seen=actual if isinstance(actual,list) else [];findings=[(isinstance(actual,list),f"{label} 应为数组"),(len(seen)==5,f"{label} 应有5道菜")]
    for index,wanted in enumerate(expected):
        row=seen[index] if index<len(seen) and isinstance(seen[index],dict) else {}
        for field in RECIPE_FIELDS:findings.append((_same(row.get(field),wanted.get(field)),f"{label} 第{index+1}道 {field} 错误"))
        actual_matched=row.get("Matched_Ingredients") if isinstance(row.get("Matched_Ingredients"),list) else []
        wanted_matched=wanted["Matched_Ingredients"]
        findings.append((len(actual_matched)==len(wanted_matched),f"{label} 第{index+1}道已报价原料数错误"))
        for item_index,wanted_item in enumerate(wanted_matched):
            actual_item=actual_matched[item_index] if item_index<len(actual_matched) and isinstance(actual_matched[item_index],dict) else {}
            for field in ("Recipe_Ingredient","Price_Item","Cheapest_Store","Quoted_Price"):
                findings.append((_same(actual_item.get(field),wanted_item[field]),f"{label} 第{index+1}道第{item_index+1}个已报价原料 {field} 错误"))
        for field in ("Unpriced_Ingredients","Cheapest_Store_Mix"):
            actual_list=row.get(field) if isinstance(row.get(field),list) else []
            wanted_list=wanted[field]
            findings.append((len(actual_list)==len(wanted_list),f"{label} 第{index+1}道 {field} 数量错误"))
            for item_index,value in enumerate(wanted_list):findings.append((item_index<len(actual_list) and actual_list[item_index]==value,f"{label} 第{index+1}道 {field}[{item_index}] 错误"))
    return findings


def _budget_findings(actual:Any,expected:dict,label:str):
    source=actual if isinstance(actual,dict) else {};findings=[(isinstance(actual,dict),f"{label} 应为对象")]
    for field in ("Weekly_Total","Average_Daily","Weekly_Budget","Budget_Headroom","Best_Single_Store","Best_Single_Store_Total","Savings_Vs_Best_Single_Store","Savings_Pct"):
        findings.append((_same(source.get(field),expected[field]),f"{label}.{field} 错误"))
    totals=source.get("Single_Store_Totals") if isinstance(source.get("Single_Store_Totals"),dict) else {}
    for store in V1.STORES:findings.append((_same(totals.get(store),expected["Single_Store_Totals"][store]),f"{label}.Single_Store_Totals.{store} 错误"))
    return findings


def _result_findings(actual:Any,expected:dict):
    source=actual if isinstance(actual,dict) else {};findings=[(isinstance(actual,dict),"结果应为对象")]
    findings.extend(_rows(source.get("price_rows"),expected["price_rows"],PRICE_FIELDS,"price_rows"))
    findings.extend(_recipe_findings(source.get("selected_recipes"),expected["selected_recipes"],"selected_recipes"))
    findings.extend(_rows(source.get("weekly_meal_plan"),expected["weekly_meal_plan"],PLAN_FIELDS,"weekly_meal_plan"))
    findings.extend(_budget_findings(source.get("budget_summary"),expected["budget_summary"],"budget_summary"))
    limitations=source.get("limitations");blob=V1.norm(" ".join(str(item) for item in limitations)) if isinstance(limitations,list) else ""
    findings.extend([(isinstance(limitations,list) and len(limitations)>=2,"limitations 少于2条"),("quoted-basket" in blob,"limitations 未披露报价篮子口径"),("cannot be verified" in blob,"limitations 未披露蛋白质无法验证")])
    return findings


def _inputs():
    root=Path(__file__).resolve().parents[1];folder=Path("/app/input_files") if Path("/app/input_files").is_dir() else root/"data"/"input_files";return folder/"grocery_prices.html",folder/"all_recipes.json"


def _core(artifact:Path):
    if not artifact.is_dir():return False,["输出目录不存在"]
    errors=[]
    for name in REQUIRED:
        path=artifact/name
        if not path.is_file() or path.stat().st_size==0:errors.append(f"核心交付缺失或为空：{name}");continue
        try:
            if name==PROCESSOR:
                text=path.read_text(encoding="utf-8");compile(text,str(path),"exec")
                if len(text.strip())<100:raise ValueError("脚本过短")
            elif name==JSON_NAME:
                if not isinstance(json.loads(path.read_text(encoding="utf-8")),dict):raise ValueError("JSON顶层不是对象")
            elif name==XLSX:
                book=load_workbook(path,read_only=True,data_only=False)
                if book.sheetnames!=["Grocery_Prices","Recipe_Cost_Analysis","Weekly_Meal_Plan"]:raise ValueError(f"工作表错误：{book.sheetnames}")
                book.close()
            elif name==DOCX:
                if not Document(path).paragraphs:raise ValueError("Word无段落")
            else:
                text=path.read_text(encoding="utf-8")
                if "BEGIN:VCALENDAR" not in text or "END:VCALENDAR" not in text:raise ValueError("ICS边界错误")
        except Exception as exc:errors.append(f"核心交付无法解析：{name}: {type(exc).__name__}: {exc}")
    return not errors,errors


def _fatal(errors,hashes):
    return {"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":hashes or ["0"*64],"criterion_results":[_criterion(key,0,"；".join(errors),0,0) for key in (*POSITIVE_WEIGHTS,*PENALTY_WEIGHTS)],"reward":0.0,"passed":False,"hurdle_failures":["R001"],"penalties":[]}


def _processor(artifact:Path,mutated:bool):
    prices,recipes=_inputs()
    with tempfile.TemporaryDirectory(prefix="cowork-evo-meal-v2-") as raw:
        folder=Path(raw)
        if mutated:
            changed=folder/"mutated-prices.html";V1.mutate_prices(prices,changed);prices=changed
        expected=V1.derive_expected(prices,recipes)
        actual,evidence=V1.run_processor(artifact/PROCESSOR,prices,recipes,folder/"result.json")
        findings=[(actual is not None,f"脚本回放失败：{evidence}")]
        if actual is not None:findings.extend(_result_findings(actual,expected))
        if mutated:
            budget=actual.get("budget_summary",{}) if isinstance(actual,dict) else {}
            findings.extend([(_same(budget.get("Weekly_Total"),expected["budget_summary"]["Weekly_Total"]),"变异周总额未传播"),(budget.get("Weekly_Total")!=7632,"变异后周总额仍是原值")])
        score,failures,earned,total=_ratio(findings);return score,failures,earned,total,actual


def _excel(artifact:Path,expected:dict):
    book=load_workbook(artifact/XLSX,data_only=False);findings=[(book.sheetnames==["Grocery_Prices","Recipe_Cost_Analysis","Weekly_Meal_Plan"],"Excel工作表名称或顺序错误")]
    price_rows,price_missing=V1.extract_rows(book["Grocery_Prices"],PRICE_FIELDS[:5]);findings.append((not price_missing,f"Grocery_Prices缺表头：{price_missing}"));findings.extend(_rows(price_rows,expected["price_rows"],PRICE_FIELDS[:5],"Excel.Grocery_Prices"))
    recipe_fields=["Recipe_ID","Recipe","Matched_Ingredient_Count","Unpriced_Ingredient_Count","Total_Cost","Cheapest_Store_Mix","Cost_Per_Serving","Unpriced_Ingredients"]
    recipe_rows,recipe_missing=V1.extract_rows(book["Recipe_Cost_Analysis"],recipe_fields);findings.append((not recipe_missing,f"Recipe_Cost_Analysis缺表头：{recipe_missing}"))
    findings.extend([(len(recipe_rows)==5,"Recipe_Cost_Analysis应有5行")])
    for index,wanted in enumerate(expected["selected_recipes"]):
        row=recipe_rows[index] if index<len(recipe_rows) else {}
        for field in ("Recipe_ID","Recipe","Matched_Ingredient_Count","Unpriced_Ingredient_Count","Total_Cost"):
            findings.append((_same(row.get(field),wanted[field]),f"Excel菜品第{index+1}行 {field} 错误"))
        findings.append((str(row.get("Cheapest_Store_Mix") or "")=="/".join(wanted["Cheapest_Store_Mix"]),f"Excel菜品第{index+1}行店铺组合错误"))
        findings.append((str(row.get("Unpriced_Ingredients") or "")=="; ".join(wanted["Unpriced_Ingredients"]),f"Excel菜品第{index+1}行未报价原料错误"))
    plan_rows,plan_missing=V1.extract_rows(book["Weekly_Meal_Plan"],PLAN_FIELDS[:5]);findings.append((not plan_missing,f"Weekly_Meal_Plan缺表头：{plan_missing}"));findings.extend(_rows(plan_rows,expected["weekly_meal_plan"],PLAN_FIELDS[:5],"Excel.Weekly_Meal_Plan"))
    book.close();score,failures,earned,total=_ratio(findings);return score,failures,earned,total,price_rows,recipe_rows,plan_rows


def _formulas(artifact:Path):
    book=load_workbook(artifact/XLSX,data_only=False);findings=[]
    prices=book["Grocery_Prices"]
    for row in range(2,24):
        f=V1.formula_norm(prices[f"F{row}"].value);g=V1.formula_norm(prices[f"G{row}"].value)
        findings.extend([(f.startswith("=") and all(token in f for token in ("INDEX(","MATCH(MIN(",f"C{row}:E{row}")),f"Grocery_Prices!F{row}最低店公式错误"),(g.startswith("=MIN(") and f"C{row}:E{row}" in g,f"Grocery_Prices!G{row}最低价公式错误")])
    recipes=book["Recipe_Cost_Analysis"]
    for row in range(2,7):
        g=V1.formula_norm(recipes[f"G{row}"].value);findings.append((g.startswith("=ROUND(") and f"E{row}/4" in g,f"Recipe_Cost_Analysis!G{row}四人份公式错误"))
    plan=book["Weekly_Meal_Plan"]
    for row in range(2,12):
        f=V1.formula_norm(plan[f"F{row}"].value);findings.append((f.startswith("=") and "RECIPE_COST_ANALYSIS!" in f and f"D{row}" in f and any(token in f for token in ("VLOOKUP(","XLOOKUP(","INDEX(")),f"Weekly_Meal_Plan!F{row}成本查找公式错误"))
    book.close();return _ratio(findings)


def _word(artifact:Path,expected:dict):
    _,text,headings,table_rows,media=V1.document_snapshot(artifact/DOCX);blob=V1.norm(text);heading_set={V1.norm(item) for item in headings};findings=[("weekly cafeteria meal proposal" in blob,"Word英文标题错误"),(table_rows>=6,"Word原生表格内容不足"),(not media,"Word用图片替代关键内容")]
    for heading in ("budget overview","menu highlights","procurement strategy","limitations"):findings.append((heading in heading_set,f"Word缺标题 {heading}"))
    for recipe in expected["selected_recipes"]:findings.append((V1.norm(recipe["Recipe"]) in blob,f"Word缺菜名 {recipe['Recipe']}"))
    for token in ("7,632","1,526.40","4,368","7,946","314","3.95","freshmart","valuegrocer","bulkbarn","quoted-basket"):
        findings.append((token in re.sub(r"\s","",text).lower() if token[0].isdigit() else token in blob,f"Word缺事实 {token}"))
    score,failures,earned,total=_ratio(findings);return score,failures,earned,total,text


def _ics(artifact:Path,expected:dict):
    events=V1.parse_ics(artifact/ICS);lunches=[row for row in expected["weekly_meal_plan"] if row["Meal"]=="Lunch"];recipes={row["Recipe_ID"]:row for row in expected["selected_recipes"]};findings=[(len(events)==5,"ICS事件数不是5")]
    for index,lunch in enumerate(lunches):
        event=events[index] if index<len(events) else {};recipe=recipes[lunch["Recipe_ID"]];stamp=lunch["Date"].replace("-","");description=event.get("DESCRIPTION","")
        findings.extend([(event.get("DTSTART;TZID=Europe/Moscow")==f"{stamp}T080000",f"事件{index+1}开始时间错误"),(event.get("DTEND;TZID=Europe/Moscow")==f"{stamp}T093000",f"事件{index+1}结束时间错误"),(event.get("SUMMARY")=="Meal Prep: "+lunch["Recipe"],f"事件{index+1}标题错误")])
        for store in recipe["Cheapest_Store_Mix"]:findings.append((store in description,f"事件{index+1}缺采购店 {store}"))
        findings.append(("Без котировки" in description,f"事件{index+1}未单列未报价原料"))
        for ingredient in recipe["Unpriced_Ingredients"]:findings.append((ingredient in description,f"事件{index+1}缺未报价原料 {ingredient}"))
    score,failures,earned,total=_ratio(findings);return score,failures,earned,total,events


def _cross(delivered:dict,excel_rows,events,word_text,expected):
    price_rows,recipe_rows,plan_rows=excel_rows;findings=[]
    json_prices={row.get("Item"):row for row in delivered.get("price_rows",[]) if isinstance(row,dict)}
    for row in price_rows:
        item=row.get("Item");findings.append((item in json_prices and all(_same(row.get(key),json_prices[item].get(key)) for key in PRICE_FIELDS[:5]),f"商品 {item} 跨JSON/Excel冲突"))
    json_recipes={row.get("Recipe_ID"):row for row in delivered.get("selected_recipes",[]) if isinstance(row,dict)}
    for row in recipe_rows:
        recipe_id=row.get("Recipe_ID");findings.append((recipe_id in json_recipes and all(_same(row.get(key),json_recipes[recipe_id].get(key)) for key in ("Recipe","Matched_Ingredient_Count","Unpriced_Ingredient_Count","Total_Cost")),f"菜品 {recipe_id} 跨JSON/Excel冲突"))
    json_plan=delivered.get("weekly_meal_plan",[])
    for index,row in enumerate(plan_rows):
        source=json_plan[index] if index<len(json_plan) and isinstance(json_plan[index],dict) else {};findings.append((all(_same(row.get(key),source.get(key)) for key in PLAN_FIELDS[:5]),f"菜单第{index+1}行跨JSON/Excel冲突"))
    lunches=[row for row in expected["weekly_meal_plan"] if row["Meal"]=="Lunch"]
    for index,lunch in enumerate(lunches):
        event=events[index] if index<len(events) else {};findings.extend([(event.get("SUMMARY")=="Meal Prep: "+lunch["Recipe"],f"第{index+1}个ICS与Lunch冲突"),(lunch["Recipe"] in word_text,f"Word未覆盖Lunch菜名 {lunch['Recipe']}")])
    return _ratio(findings)


def _formatting(artifact:Path):
    book=load_workbook(artifact/XLSX,data_only=False);details=V1.workbook_formatting(book);findings=[]
    for sheet,item in details.items():findings.extend([(item.get("tables",0)>=1,f"{sheet}没有原生Excel表"),(item.get("freeze")=="A2",f"{sheet}未冻结首行"),(item.get("header_bold"),f"{sheet}表头未全部加粗"),(item.get("header_fill"),f"{sheet}表头未全部填色"),(item.get("wide_columns"),f"{sheet}列宽不足")])
    book.close();document=Document(artifact/DOCX);findings.extend([(len(document.tables)>=2,"Word原生表少于2张"),(sum(paragraph.style.name.startswith("Heading") for paragraph in document.paragraphs)>=5,"Word标题层级不足")]);return _ratio(findings)


def _render(artifact:Path):
    with tempfile.TemporaryDirectory(prefix="cowork-evo-meal-render-") as raw:recalculated,texts,pages,evidence=V1.libreoffice_prepare(artifact/XLSX,artifact/DOCX,Path(raw))
    x=V1.norm(texts.get(XLSX,""));d=V1.norm(texts.get(DOCX,""));findings=[(recalculated is not None,f"Excel未重算：{evidence}"),(pages.get(XLSX)==3,f"Excel渲染页数错误：{pages.get(XLSX)}"),(pages.get(DOCX,0)>=2,f"Word渲染页数不足：{pages.get(DOCX)}"),(all(token in x for token in ("item","cheapest_store","recipe_id","cost_per_serving","estimated_cost","гречка с тушёнкой")),"Excel渲染缺核心文字"),(all(token in d for token in ("weekly cafeteria meal proposal","budget overview","procurement strategy","limitations")),"Word渲染缺核心文字")];return _ratio(findings)


def _limits(delivered:dict,word_text:str):
    limitations=delivered.get("limitations");json_blob=V1.norm(" ".join(str(item) for item in limitations)) if isinstance(limitations,list) else "";word_blob=V1.norm(word_text);return _ratio([("quoted-basket" in json_blob,"JSON未披露报价篮子口径"),("cannot be verified" in json_blob,"JSON未披露蛋白质不可验证"),("quoted-basket" in word_blob,"Word未披露报价篮子口径"),("не импутировал" in word_blob,"Word未披露未报价不估值"),("невозможно проверить" in word_blob,"Word未披露蛋白质无法验证")])


def _selection_penalty(delivered:dict):
    _,recipes_path=_inputs();recipes=json.loads(recipes_path.read_text(encoding="utf-8"));catalog={row["id"]:row for row in recipes};prices,_=_inputs();truth=V1.derive_expected(prices,recipes_path);eligible_ids=set()
    price_items={row["Item"] for row in truth["price_rows"]}
    for recipe in recipes:
        matched={V1.ALIASES.get(str(item.get("name") or "").strip(),str(item.get("name") or "").strip()) for item in recipe.get("ingredients",[])}&price_items
        if recipe.get("category")=="горячее" and len(matched)>=3:eligible_ids.add(recipe["id"])
    selected=[row.get("Recipe_ID") for row in delivered.get("selected_recipes",[]) if isinstance(row,dict)];invalid=[item for item in selected if item not in eligible_ids or item not in catalog];severity=min(1,len(invalid)/2) if invalid else 0;return severity,([f"入选非热菜或少于3个已报价原料：{invalid}"] if invalid else []),1.0


def _script_penalty(canonical_score,mutated_score,canonical,mutated):
    labels=[];severity=0.0
    if canonical is None:severity=max(severity,.65);labels.append("原始输入回放失败")
    if mutated is None:severity=max(severity,.8);labels.append("变异价格输入回放失败")
    elif canonical is not None and canonical==mutated:severity=1.0;labels.append("脚本对价格变异完全无响应")
    elif canonical_score>.9 and mutated_score<.5:severity=max(severity,.75);labels.append("脚本疑似写死原始结果")
    return severity,labels,1.0


def _extra_penalty(artifact:Path):
    allowed={*REQUIRED,"solve.py","solve.sh","__pycache__"};extras=sorted(path.name for path in artifact.iterdir() if path.name not in allowed);return (1.0 if extras else 0.0),([f"输出含额外文件或目录：{extras}"] if extras else []),1.0


def _unsupported(delivered:dict,word_text:str):
    findings=V1.unsupported_claims(json.dumps(delivered,ensure_ascii=False)+"\n"+word_text);severity=min(1,.6+.2*len(findings)) if findings else 0;return severity,findings,1.0


def _safe(call,*args):
    try:return call(*args)
    except Exception as exc:return 0.0,[f"任务专属检查异常：{type(exc).__name__}: {exc}"],0,1


def _safe_penalty(call,*args):
    try:return call(*args)
    except Exception as exc:return 1.0,[f"严重边界无法核验：{type(exc).__name__}: {exc}"],1.0


def evaluate(artifact:Path):
    ok,errors=_core(artifact);hashes=[_sha(artifact/name) for name in REQUIRED if (artifact/name).is_file()]
    if not ok:return _fatal(errors,hashes)
    prices,recipes=_inputs();expected=V1.derive_expected(prices,recipes);delivered=json.loads((artifact/JSON_NAME).read_text(encoding="utf-8"));results=[_criterion("R001",1,"五份核心交付存在、非空且可解析",5,5)];scores={"R001":1.0}
    canonical_score,canonical_failures,canonical_earned,canonical_total,canonical=_processor(artifact,False);mutated_score,mutated_failures,mutated_earned,mutated_total,mutated=_processor(artifact,True)
    for key,score,failures,earned,total in (("R002",canonical_score,canonical_failures,canonical_earned,canonical_total),("R003",mutated_score,mutated_failures,mutated_earned,mutated_total)):
        results.append(_criterion(key,score,"全部任务专属原子检查通过" if not failures else "未通过："+"；".join(failures),earned,total));scores[key]=score
    groups={"R004":_ratio(_rows(delivered.get("price_rows"),expected["price_rows"],PRICE_FIELDS,"JSON.price_rows")),"R005":_ratio(_recipe_findings(delivered.get("selected_recipes"),expected["selected_recipes"],"JSON.selected_recipes")),"R006":_ratio(_rows(delivered.get("weekly_meal_plan"),expected["weekly_meal_plan"],PLAN_FIELDS,"JSON.weekly_meal_plan")+_budget_findings(delivered.get("budget_summary"),expected["budget_summary"],"JSON.budget_summary"))}
    excel_score,excel_failures,excel_earned,excel_total,*excel_rows=_excel(artifact,expected);groups["R007"]=(excel_score,excel_failures,excel_earned,excel_total);groups["R008"]=_formulas(artifact)
    word_score,word_failures,word_earned,word_total,word_text=_word(artifact,expected);groups["R009"]=(word_score,word_failures,word_earned,word_total)
    ics_score,ics_failures,ics_earned,ics_total,events=_ics(artifact,expected);groups["R010"]=(ics_score,ics_failures,ics_earned,ics_total)
    groups["R011"]=_cross(delivered,excel_rows,events,word_text,expected);groups["R012"]=_formatting(artifact);groups["R013"]=_render(artifact);groups["R014"]=_limits(delivered,word_text)
    for key in [f"R{index:03d}" for index in range(4,15)]:
        score,failures,earned,total=groups[key];results.append(_criterion(key,score,"全部任务专属原子检查通过" if not failures else "未通过："+"；".join(failures),earned,total));scores[key]=score
    penalties={"R015":_safe_penalty(_selection_penalty,delivered),"R016":_safe_penalty(_script_penalty,canonical_score,mutated_score,canonical,mutated),"R017":_safe_penalty(_extra_penalty,artifact),"R018":_safe_penalty(_unsupported,delivered,word_text)}
    for key,(score,labels,total) in penalties.items():results.append(_criterion(key,score,"未触发任务专属严重边界" if not labels else "；".join(labels),score*total,total));scores[key]=score
    total=sum(POSITIVE_WEIGHTS.values());positive=sum(POSITIVE_WEIGHTS[key]*scores[key] for key in POSITIVE_WEIGHTS);penalty=sum(abs(PENALTY_WEIGHTS[key])*scores[key] for key in PENALTY_WEIGHTS);reward=round(max(0,min(1,(positive-penalty)/total)),6)
    return {"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":hashes,"criterion_results":results,"reward":reward,"passed":reward>=.8,"hurdle_failures":[],"penalties":[{"id":key,"amount":round(abs(PENALTY_WEIGHTS[key])*scores[key]/total,6)} for key in PENALTY_WEIGHTS if scores[key]>0]}


def main():
    parser=argparse.ArgumentParser();parser.add_argument("--artifact-dir",type=Path,required=True);parser.add_argument("--result",type=Path,required=True);parser.add_argument("--reward-file",type=Path,required=True);parser.add_argument("--scoring-version",choices=("v1","v2"),default="v2");args=parser.parse_args()
    try:result=V1.evaluate(args.artifact_dir) if args.scoring_version=="v1" else evaluate(args.artifact_dir)
    except Exception as exc:result=_fatal([f"Judge异常：{type(exc).__name__}: {exc}"],[])
    rendered=json.dumps(result,ensure_ascii=False,indent=2);args.result.parent.mkdir(parents=True,exist_ok=True);args.reward_file.parent.mkdir(parents=True,exist_ok=True);args.result.write_text(rendered+"\n",encoding="utf-8");args.reward_file.write_text(f"{float(result['reward']):.6f}\n",encoding="utf-8");print(rendered)


if __name__=="__main__":main()
