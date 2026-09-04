#!/usr/bin/env python3
"""食堂菜单任务的任务专属 v2 致命/局部反例。"""
from __future__ import annotations
import importlib.util,json,shutil,sys,tempfile
from pathlib import Path
from docx import Document
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1];GOLDEN=ROOT/"solution"
REQUIRED=("meal_cost_optimizer.py","meal_optimization_results.json","Weekly_Meal_Budget.xlsx","Cafeteria_Meal_Proposal.docx","meal_prep_calendar.ics")

def _load():
    path=Path(__file__).with_name("judge.py");spec=importlib.util.spec_from_file_location("cafeteria_v2_negative_judge",path)
    if spec is None or spec.loader is None:raise RuntimeError("无法载入Judge")
    module=importlib.util.module_from_spec(spec);sys.modules[spec.name]=module;spec.loader.exec_module(module);return module
JUDGE=_load()

def copy(root:Path):
    root.mkdir(parents=True,exist_ok=True)
    for name in REQUIRED:shutil.copy2(GOLDEN/name,root/name)
def load_json(root):return json.loads((root/"meal_optimization_results.json").read_text(encoding="utf-8"))
def save_json(root,value):(root/"meal_optimization_results.json").write_text(json.dumps(value,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
def json_price(root):
    value=load_json(root);value["price_rows"][0]["FreshMart_Price"]+=7;save_json(root,value)
def json_store(root):
    value=load_json(root);value["price_rows"][0]["Cheapest_Store"]="FreshMart";save_json(root,value)
def json_recipe_cost(root):
    value=load_json(root);value["selected_recipes"][0]["Total_Cost"]+=11;save_json(root,value)
def json_ingredient(root):
    value=load_json(root);value["selected_recipes"][0]["Matched_Ingredients"][0]["Quoted_Price"]+=4;save_json(root,value)
def json_recipe_id(root):
    value=load_json(root);value["selected_recipes"][0]["Recipe_ID"]="salat-salat-olive";save_json(root,value)
def json_plan_date(root):
    value=load_json(root);value["weekly_meal_plan"][0]["Date"]="2026-03-15";save_json(root,value)
def json_plan_cost(root):
    value=load_json(root);value["weekly_meal_plan"][0]["Estimated_Cost"]+=10;save_json(root,value)
def json_budget(root):
    value=load_json(root);value["budget_summary"]["Weekly_Total"]+=10;save_json(root,value)
def json_best_store(root):
    value=load_json(root);value["budget_summary"]["Best_Single_Store"]="BulkBarn";save_json(root,value)
def processor_category(root):
    path=root/"meal_cost_optimizer.py";text=path.read_text(encoding="utf-8");path.write_text(text.replace('recipe.get("category") != "горячее"','recipe.get("category") != "суп"'),encoding="utf-8")
def processor_hardcoded(root):
    payload=load_json(root);script="#!/usr/bin/env python3\nimport argparse,json\nfrom pathlib import Path\np=argparse.ArgumentParser();p.add_argument('--prices');p.add_argument('--recipes');p.add_argument('--output',type=Path);a=p.parse_args()\npayload="+repr(payload)+"\na.output.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+'\\n',encoding='utf-8')\n";(root/"meal_cost_optimizer.py").write_text(script,encoding="utf-8")
def excel_cell(root,sheet,cell,value):
    path=root/"Weekly_Meal_Budget.xlsx";book=load_workbook(path);book[sheet][cell]=value;book.save(path);book.close()
def excel_price(root):excel_cell(root,"Grocery_Prices","C2",999)
def excel_recipe(root):excel_cell(root,"Recipe_Cost_Analysis","E2",598)
def excel_formula(root):excel_cell(root,"Recipe_Cost_Analysis","G2",146.75)
def excel_plan(root):excel_cell(root,"Weekly_Meal_Plan","D2","goryachee-golubtsy")
def ics_start(root):
    path=root/"meal_prep_calendar.ics";text=path.read_text(encoding="utf-8");path.write_text(text.replace("20260316T080000","20260316T083000",1),encoding="utf-8")
def word_heading(root):
    path=root/"Cafeteria_Meal_Proposal.docx";doc=Document(path)
    for paragraph in doc.paragraphs:
        if paragraph.text.strip()=="Budget Overview":paragraph.text="Budget";break
    doc.save(path)
def word_protein_claim(root):
    path=root/"Cafeteria_Meal_Proposal.docx";doc=Document(path);doc.add_paragraph("Каждое блюдо обеспечивает минимум 20 г белка.");doc.save(path)
def extra(root):(root/"debug-menu.txt").write_text("temporary\n",encoding="utf-8")

PARTIAL={
 "json_one_price_wrong":json_price,"json_one_cheapest_store_wrong":json_store,"json_one_recipe_total_wrong":json_recipe_cost,
 "json_one_matched_ingredient_wrong":json_ingredient,"json_ineligible_recipe_selected":json_recipe_id,"json_one_plan_date_wrong":json_plan_date,
 "json_one_plan_cost_wrong":json_plan_cost,"json_weekly_total_wrong":json_budget,"json_best_store_wrong":json_best_store,
 "processor_hot_category_rule_wrong":processor_category,"processor_hardcoded_against_price_mutation":processor_hardcoded,
 "excel_one_source_price_wrong":excel_price,"excel_one_recipe_total_wrong":excel_recipe,"excel_cost_per_serving_formula_flattened":excel_formula,
 "excel_one_lunch_recipe_wrong":excel_plan,"ics_one_start_time_wrong":ics_start,"word_required_heading_missing":word_heading,
 "word_unsupported_protein_claim":word_protein_claim,"unexpected_debug_file":extra,
}

def main():
    render_score,render_failures,_,_=JUDGE._render(GOLDEN);assert render_score==1,render_failures
    JUDGE._render=lambda _artifact:(1.0,[],5,5)
    records=[]
    with tempfile.TemporaryDirectory(prefix="cowork-evo-meal-neg-") as raw:
        base=Path(raw)
        for name in REQUIRED:
            for mode in ("missing","empty","corrupt"):
                case=base/f"fatal-{mode}-{name}";copy(case);target=case/name
                if mode=="missing":target.unlink()
                elif mode=="empty":target.write_bytes(b"")
                else:target.write_bytes(b"not a valid task artifact")
                reward=JUDGE.evaluate(case)["reward"];assert reward==0,(mode,name,reward);records.append({"case":f"{mode}:{name}","class":"fatal","reward":reward})
        case=base/"fatal-empty-output";case.mkdir();reward=JUDGE.evaluate(case)["reward"];assert reward==0;records.append({"case":"empty-output","class":"fatal","reward":reward})
        case=base/"fatal-input-only";shutil.copytree(ROOT/"data"/"input_files",case);reward=JUDGE.evaluate(case)["reward"];assert reward==0;records.append({"case":"input-only","class":"fatal","reward":reward})
        for name,mutate in PARTIAL.items():
            case=base/f"partial-{name}";copy(case);mutate(case);reward=JUDGE.evaluate(case)["reward"];assert 0<reward<1,(name,reward);records.append({"case":name,"class":"partial","reward":reward})
    fatal=[row["reward"] for row in records if row["class"]=="fatal"];partial=[row["reward"] for row in records if row["class"]=="partial"]
    print(json.dumps({"fatal_count":len(fatal),"fatal_unique_rewards":sorted(set(fatal)),"partial_count":len(partial),"partial_min":min(partial),"partial_max":max(partial),"cases":records},ensure_ascii=False,indent=2))

if __name__=="__main__":main()
