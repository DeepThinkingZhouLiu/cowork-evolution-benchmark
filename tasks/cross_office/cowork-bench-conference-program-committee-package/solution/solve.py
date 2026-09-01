#!/usr/bin/env python3
"""逐题生成 ДРВБ-2026 程序委员会任务的四个 Golden 交付物。"""

from __future__ import annotations

import argparse
import json
import mailbox
import re
import shutil
import subprocess
import tempfile
import zipfile
from datetime import date, datetime, time, timedelta, timezone
from email import policy
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


XLSX_NAME = "Submissions.xlsx"
DOCX_NAME = "Conference_Program.docx"
PAGE_NAME = "Program_Page.md"
MBOX_NAME = "Author_Notifications.mbox"
FIXED_DT = datetime(2026, 3, 2, 9, 0, 0, tzinfo=timezone.utc)

NAVY = "18324A"
TEAL = "0E7C86"
GOLD = "C89B3C"
WHITE = "FFFFFF"
PALE_BLUE = "EAF1F8"
PALE_GREEN = "E3F2E8"
PALE_RED = "FCE8E6"
PALE_GOLD = "FFF4D6"
LINE = "C9D5DF"
GRAY = "586675"

SECTION_PLAN = {
    "Археология": (date(2026, 3, 17), time(10, 0)),
    "Источниковедение": (date(2026, 3, 18), time(10, 0)),
    "История": (date(2026, 3, 19), time(10, 0)),
}


def read_sources(input_dir: Path) -> list[dict[str, object]]:
    submissions_path = input_dir / "conference_submissions.json"
    scores_path = input_dir / "review_scores.md"
    rules_path = input_dir / "conference_rules.md"
    for path in (submissions_path, scores_path, rules_path):
        if not path.is_file():
            raise FileNotFoundError(path)

    payload = json.loads(submissions_path.read_text(encoding="utf-8"))
    responses = payload.get("responses")
    if not isinstance(responses, list) or len(responses) != 10:
        raise ValueError("conference_submissions.json 必须恰好含 10 份申请")

    score_map: dict[int, tuple[int, int, int]] = {}
    for line in scores_path.read_text(encoding="utf-8").splitlines():
        match = re.match(r"^\|\s*(\d+)\s*\|.*?\|\s*(\d)\s*\|\s*(\d)\s*\|\s*(\d)\s*\|$", line)
        if match:
            score_map[int(match.group(1))] = tuple(int(match.group(index)) for index in (2, 3, 4))
    if set(score_map) != set(range(1, 11)):
        raise ValueError(f"review_scores.md 评分 ID 不完整：{sorted(score_map)}")

    rules_text = rules_path.read_text(encoding="utf-8")
    for required in ("2026-03-17", "2026-03-18", "2026-03-19", "平均分大于或等于 4.0", "每个报告 30 分钟"):
        if required not in rules_text:
            raise ValueError(f"conference_rules.md 缺少规则：{required}")

    rows: list[dict[str, object]] = []
    required_fields = ["ФИО автора", "Email", "Аффилиация", "Название доклада", "Тематика", "Аннотация"]
    for response in sorted(responses, key=lambda item: int(item["id"])):
        identifier = int(response["id"])
        if any(not str(response.get(field, "")).strip() for field in required_fields):
            raise ValueError(f"申请 {identifier} 缺字段")
        scores = score_map[identifier]
        average = round(sum(scores) / 3, 2)
        accepted = average >= 4.0
        rows.append(
            {
                "id": identifier,
                "author": str(response["ФИО автора"]),
                "email": str(response["Email"]),
                "affiliation": str(response["Аффилиация"]),
                "title": str(response["Название доклада"]),
                "topic": str(response["Тематика"]),
                "abstract": str(response["Аннотация"]),
                "scores": scores,
                "average": average,
                "decision": "Принят" if accepted else "Отклонён",
                "section": str(response["Тематика"]) if accepted else "",
            }
        )
    return rows


def build_schedule(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    schedule: list[dict[str, object]] = []
    for section, (event_date, first_start) in SECTION_PLAN.items():
        section_rows = [row for row in rows if row["decision"] == "Принят" and row["section"] == section]
        if not 2 <= len(section_rows) <= 4:
            raise ValueError(f"分会场人数不满足 2–4：{section}={len(section_rows)}")
        for offset, row in enumerate(sorted(section_rows, key=lambda item: int(item["id"]))):
            start_dt = datetime.combine(event_date, first_start) + timedelta(minutes=30 * offset)
            end_dt = start_dt + timedelta(minutes=30)
            schedule.append({**row, "date": event_date, "start": start_dt.time(), "end": end_dt.time()})
    return sorted(schedule, key=lambda item: (item["date"], item["start"], item["id"]))


def style_sheet(sheet, widths: dict[str, float], paper_size: int | None = None) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 34
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", size=8, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color=LINE)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=8, color=NAVY)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = paper_size or sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:1"
    sheet.page_margins.left = 0.2
    sheet.page_margins.right = 0.2
    sheet.page_margins.top = 0.35
    sheet.page_margins.bottom = 0.35


def add_native_table(sheet, name: str) -> None:
    table = Table(displayName=name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def build_workbook(rows: list[dict[str, object]], schedule: list[dict[str, object]], output_path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.properties.creator = "EvoCowork benchmark adapter"
    workbook.properties.title = "DRVB-2026 Program Committee Package"
    workbook.properties.created = FIXED_DT.replace(tzinfo=None)
    workbook.properties.modified = FIXED_DT.replace(tzinfo=None)

    all_sheet = workbook.create_sheet("All_Submissions")
    all_headers = [
        "ID", "ФИО автора", "Email", "Аффилиация", "Название доклада", "Тематика", "Аннотация",
        "Рецензент 1", "Рецензент 2", "Рецензент 3", "Средняя оценка", "Статус", "Секция",
    ]
    all_sheet.append(all_headers)
    for excel_row, row in enumerate(rows, start=2):
        scores = row["scores"]
        all_sheet.append(
            [
                row["id"], row["author"], row["email"], row["affiliation"], row["title"], row["topic"], row["abstract"],
                scores[0], scores[1], scores[2],
                f"=ROUND(AVERAGE(H{excel_row}:J{excel_row}),2)",
                f'=IF(K{excel_row}>=4,"Принят","Отклонён")',
                f'=IF(L{excel_row}="Принят",F{excel_row},"")',
            ]
        )
    style_sheet(
        all_sheet,
        {"A": 5, "B": 20, "C": 25, "D": 30, "E": 45, "F": 22, "G": 52,
         "H": 12, "I": 12, "J": 12, "K": 15, "L": 14, "M": 22},
        paper_size=all_sheet.PAPERSIZE_A3,
    )
    for row in range(2, all_sheet.max_row + 1):
        all_sheet.row_dimensions[row].height = 58
        all_sheet.cell(row, 11).number_format = "0.00"
    score_validation = DataValidation(type="whole", operator="between", formula1="0", formula2="5", allow_blank=False)
    all_sheet.add_data_validation(score_validation)
    score_validation.add(f"H2:J{all_sheet.max_row}")
    all_sheet.conditional_formatting.add(
        f"L2:M{all_sheet.max_row}", FormulaRule(formula=['$L2="Принят"'], fill=PatternFill("solid", fgColor=PALE_GREEN))
    )
    all_sheet.conditional_formatting.add(
        f"L2:M{all_sheet.max_row}", FormulaRule(formula=['$L2="Отклонён"'], fill=PatternFill("solid", fgColor=PALE_RED))
    )
    add_native_table(all_sheet, "AllSubmissionsTable")
    all_sheet.print_area = f"A1:M{all_sheet.max_row}"

    program = workbook.create_sheet("Accepted_Program")
    program_headers = ["Секция", "День", "Начало", "Окончание", "Докладчик", "Email", "Название доклада", "Аффилиация", "Source ID"]
    program.append(program_headers)
    source_excel_rows = {int(row["id"]): int(row["id"]) + 1 for row in rows}
    for excel_row, slot in enumerate(schedule, start=2):
        source_row = source_excel_rows[int(slot["id"])]
        program.append(
            [
                f"='All_Submissions'!M{source_row}",
                datetime.combine(slot["date"], time.min),
                slot["start"],
                f"=C{excel_row}+TIME(0,30,0)",
                f"='All_Submissions'!B{source_row}",
                f"='All_Submissions'!C{source_row}",
                f"='All_Submissions'!E{source_row}",
                f"='All_Submissions'!D{source_row}",
                f"='All_Submissions'!A{source_row}",
            ]
        )
    style_sheet(program, {"A": 23, "B": 14, "C": 11, "D": 11, "E": 21, "F": 26, "G": 50, "H": 32, "I": 10})
    for row in range(2, program.max_row + 1):
        program.row_dimensions[row].height = 45
        program.cell(row, 2).number_format = "yyyy-mm-dd"
        program.cell(row, 3).number_format = "hh:mm"
        program.cell(row, 4).number_format = "hh:mm"
    program.conditional_formatting.add(
        f"A2:I{program.max_row}", FormulaRule(formula=['$A2="Археология"'], fill=PatternFill("solid", fgColor=PALE_GOLD))
    )
    add_native_table(program, "AcceptedProgramTable")
    program.print_area = f"A1:I{program.max_row}"

    summary = workbook.create_sheet("Control_Summary")
    summary.append(["Metric", "Value"])
    metrics = [
        ("Total_Submissions", "=COUNTA('All_Submissions'!A2:A11)"),
        ("Accepted", '=COUNTIF(\'All_Submissions\'!L2:L11,"Принят")'),
        ("Rejected", '=COUNTIF(\'All_Submissions\'!L2:L11,"Отклонён")'),
        ("Археология", '=COUNTIF(\'Accepted_Program\'!A2:A8,A5)'),
        ("Источниковедение", '=COUNTIF(\'Accepted_Program\'!A2:A8,A6)'),
        ("История", '=COUNTIF(\'Accepted_Program\'!A2:A8,A7)'),
        ("Earliest_Start", "=MIN('Accepted_Program'!C2:C8)"),
        ("Latest_End", "=MAX('Accepted_Program'!D2:D8)"),
        (
            "Schedule_Conflicts",
            "=IF(OR(AND('Accepted_Program'!B2='Accepted_Program'!B3,'Accepted_Program'!D2>'Accepted_Program'!C3),"
            "AND('Accepted_Program'!B3='Accepted_Program'!B4,'Accepted_Program'!D3>'Accepted_Program'!C4),"
            "AND('Accepted_Program'!B4='Accepted_Program'!B5,'Accepted_Program'!D4>'Accepted_Program'!C5),"
            "AND('Accepted_Program'!B5='Accepted_Program'!B6,'Accepted_Program'!D5>'Accepted_Program'!C6),"
            "AND('Accepted_Program'!B6='Accepted_Program'!B7,'Accepted_Program'!D6>'Accepted_Program'!C7),"
            "AND('Accepted_Program'!B7='Accepted_Program'!B8,'Accepted_Program'!D7>'Accepted_Program'!C8)),1,0)",
        ),
    ]
    for label, formula in metrics:
        summary.append([label, formula])
    style_sheet(summary, {"A": 34, "B": 24})
    summary.page_setup.orientation = "portrait"
    for row in range(2, summary.max_row + 1):
        summary.cell(row, 1).fill = PatternFill("solid", fgColor=PALE_BLUE)
        summary.cell(row, 1).font = Font(name="Aptos", size=10, bold=True, color=NAVY)
        summary.cell(row, 2).font = Font(name="Aptos Display", size=12, bold=True, color=TEAL)
    summary.cell(8, 2).number_format = "hh:mm"
    summary.cell(9, 2).number_format = "hh:mm"
    summary.conditional_formatting.add("B10", FormulaRule(formula=["B10>0"], fill=PatternFill("solid", fgColor=PALE_RED)))
    add_native_table(summary, "ControlSummaryTable")
    summary.print_area = f"A1:B{summary.max_row}"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    recalculate_workbook(output_path)
    restore_frozen_header_rows(output_path)
    normalize_ooxml(output_path)


def recalculate_workbook(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="conference-xlsx-recalc-") as tmp:
        tmp_root = Path(tmp)
        input_root = tmp_root / "input"
        output_root = tmp_root / "output"
        profile = tmp_root / "profile"
        input_root.mkdir()
        output_root.mkdir()
        profile.mkdir()
        source = input_root / path.name
        shutil.copy2(path, source)
        command = [
            "soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "xlsx",
            "--outdir", str(output_root), str(source),
        ]
        result = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300)
        converted = output_root / path.name
        if result.returncode != 0 or not converted.is_file():
            raise RuntimeError(f"LibreOffice 公式重算失败：{result.stdout[-1000:]}")
        shutil.copy2(converted, path)


def restore_frozen_header_rows(path: Path) -> None:
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with tempfile.TemporaryDirectory(prefix="conference-xlsx-freeze-") as tmp:
        target = Path(tmp) / path.name
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w") as destination:
            for info in source.infolist():
                data = source.read(info.filename)
                if re.match(r"xl/worksheets/sheet[1-3]\.xml$", info.filename):
                    root = etree.fromstring(data)
                    sheet_view = root.find(f"{{{namespace}}}sheetViews/{{{namespace}}}sheetView")
                    if sheet_view is not None:
                        for existing in sheet_view.findall(f"{{{namespace}}}pane"):
                            sheet_view.remove(existing)
                        pane = etree.Element(f"{{{namespace}}}pane")
                        pane.set("ySplit", "1")
                        pane.set("topLeftCell", "A2")
                        pane.set("activePane", "bottomLeft")
                        pane.set("state", "frozen")
                        sheet_view.insert(0, pane)
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                destination.writestr(info, data)
        shutil.copy2(target, path)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE, size: str = "6") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def add_word_table(document: Document, headers: list[str], data: list[list[str]], widths: list[float]):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, NAVY)
        set_cell_border(cell, NAVY)
        cell.width = Cm(widths[index])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(8)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(data):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.text = str(value)
            set_cell_shading(cell, WHITE if row_index % 2 == 0 else "F4F7FA")
            set_cell_border(cell)
            cell.width = Cm(widths[index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(8)
                    run.font.color.rgb = RGBColor.from_string(NAVY)
    return table


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("ДРВБ-2026  •  ")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)


def build_program_document(schedule: list[dict[str, object]], output_path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    document.core_properties.creator = "EvoCowork benchmark adapter"
    document.core_properties.title = "Программа конференции ДРВБ-2026"
    document.core_properties.created = FIXED_DT.replace(tzinfo=None)
    document.core_properties.modified = FIXED_DT.replace(tzinfo=None)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].font.color.rgb = RGBColor.from_string(NAVY)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for name, size, color in (("Title", 27, NAVY), ("Subtitle", 13, TEAL), ("Heading 1", 18, NAVY), ("Heading 2", 12, TEAL)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = "КАЗАНСКИЙ ФЕДЕРАЛЬНЫЙ УНИВЕРСИТЕТ  /  17–19 МАРТА 2026"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(TEAL)
    add_page_number(section.footer.paragraphs[0])

    title = document.add_paragraph(style="Title")
    title.add_run("Древняя Русь и\nВолжская Булгария")
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run("история и археология • памяти А.П. Смирнова")
    document.add_paragraph("Официальная программа конференции", style="Heading 2")
    document.add_paragraph("17–19 марта 2026 года")
    document.add_paragraph("Казанский федеральный университет (КФУ)\nИнститут международных отношений")
    callout = document.add_table(rows=1, cols=3)
    callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, heading, value in zip(
        callout.rows[0].cells,
        ("ДНЕЙ", "СЕКЦИЙ", "ДОКЛАДОВ"),
        ("3", "3", "7"),
        strict=True,
    ):
        set_cell_shading(cell, PALE_BLUE)
        set_cell_border(cell, TEAL)
        cell.text = f"{heading}\n{value}"
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.name = "Aptos Display"
                run.font.bold = True
                run.font.size = Pt(11)
                run.font.color.rgb = RGBColor.from_string(NAVY)
    document.add_paragraph(
        "Конференция посвящена памяти А.П. Смирнова и объединяет исследования по истории и археологии "
        "Древней Руси и Волжской Булгарии. Программа сформирована по итогам рецензирования заявок."
    )

    document.add_page_break()
    month_names = {17: "17 марта 2026", 18: "18 марта 2026", 19: "19 марта 2026"}
    for section_name, (event_date, _) in SECTION_PLAN.items():
        document.add_heading(month_names[event_date.day], level=1)
        document.add_paragraph(section_name, style="Subtitle")
        section_slots = [slot for slot in schedule if slot["section"] == section_name]
        data = [
            [
                f"{slot['start'].strftime('%H:%M')}–{slot['end'].strftime('%H:%M')}",
                f"{slot['author']}\n{slot['affiliation']}",
                str(slot["title"]),
            ]
            for slot in section_slots
        ]
        add_word_table(document, ["Время", "Докладчик", "Название доклада"], data, [2.8, 5.3, 9.0])
        document.add_paragraph("Все временные интервалы включают доклад и обсуждение.")

    document.add_heading("Площадка", level=1)
    document.add_paragraph("Казанский федеральный университет (КФУ), Институт международных отношений.")
    document.add_heading("О конференции", level=1)
    document.add_paragraph(
        "Конференция «Древняя Русь и Волжская Булгария: история и археология» проводится в память А.П. Смирнова. "
        "Настоящий документ содержит только сведения из утверждённой программы; дополнительные биографические "
        "данные и организационные детали не включены."
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    normalize_ooxml(output_path)


def build_program_page(schedule: list[dict[str, object]], output_path: Path) -> None:
    lines = [
        "# Программа конференции ДРВБ-2026",
        "",
        "**Древняя Русь и Волжская Булгария: история и археология. Памяти А.П. Смирнова**",
        "",
        "17–19 марта 2026 · Казанский федеральный университет (КФУ) · Институт международных отношений",
    ]
    month_names = {17: "17 марта 2026", 18: "18 марта 2026", 19: "19 марта 2026"}
    for section_name, (event_date, _) in SECTION_PLAN.items():
        lines.extend(["", f"## {month_names[event_date.day]} · {section_name}", "", "| Время       | Докладчик | Название доклада | Аффилиация |", "|-------------|-----------|------------------|------------|"])
        for slot in [item for item in schedule if item["section"] == section_name]:
            lines.append(
                f"| {slot['start'].strftime('%H:%M')}–{slot['end'].strftime('%H:%M')} | {slot['author']} | "
                f"{str(slot['title']).replace('|', '—')} | {str(slot['affiliation']).replace('|', '—')} |"
            )
    lines.extend(["", "Каждый 30-минутный интервал включает доклад и обсуждение.", ""])
    output_path.write_text("\n".join(lines), encoding="utf-8", newline="\n")


def build_notifications(rows: list[dict[str, object]], schedule: list[dict[str, object]], output_path: Path) -> None:
    schedule_by_id = {int(item["id"]): item for item in schedule}
    chunks: list[bytes] = []
    for index, row in enumerate(rows, start=1):
        message = EmailMessage(policy=policy.default.clone(linesep="\n"))
        message["Date"] = format_datetime(FIXED_DT + timedelta(minutes=index))
        message["From"] = "program-committee@drvb2026.local"
        message["To"] = str(row["email"])
        message["Message-ID"] = f"<drvb2026-{int(row['id']):02d}@drvb2026.local>"
        surname = str(row["author"]).split()[0]
        if row["decision"] == "Принят":
            slot = schedule_by_id[int(row["id"])]
            message["Subject"] = "Ваш доклад принят на ДРВБ-2026"
            body = (
                f"Уважаемый(ая) {surname},\n\n"
                f"Ваш доклад «{row['title']}» принят в программу ДРВБ-2026.\n"
                f"Секция: {slot['section']}.\n"
                f"Дата и время начала: {slot['date'].isoformat()}, {slot['start'].strftime('%H:%M')}.\n\n"
                "Конференция состоится 17–19 марта 2026 в Казанском федеральном университете.\n\n"
                "С уважением,\nПрограммный комитет ДРВБ-2026"
            )
        else:
            message["Subject"] = "Решение по заявке ДРВБ-2026"
            body = (
                f"Уважаемый(ая) {surname},\n\n"
                f"Благодарим за заявку с докладом «{row['title']}». К сожалению, доклад не включён в программу ДРВБ-2026.\n"
                "Конференция состоится 17–19 марта 2026. Благодарим за интерес и желаем успехов в дальнейшей работе.\n\n"
                "С уважением,\nПрограммный комитет ДРВБ-2026"
            )
        message.set_content(body, subtype="plain", charset="utf-8")
        separator = f"From program-committee@drvb2026.local Mon Mar 02 09:{index:02d}:00 2026\n".encode("ascii")
        chunks.append(separator + message.as_bytes() + b"\n")
    output_path.write_bytes(b"".join(chunks))
    parsed = mailbox.mbox(output_path, create=False)
    if len(parsed) != 10:
        raise RuntimeError(f"生成的 mbox 无法解析为 10 封邮件：{len(parsed)}")
    parsed.close()


def normalize_ooxml(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="conference-ooxml-normalize-") as tmp:
        target = Path(tmp) / path.name
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w") as destination:
            for info in sorted(source.infolist(), key=lambda item: item.filename):
                data = source.read(info.filename)
                if info.filename == "docProps/core.xml":
                    root = etree.fromstring(data)
                    dcterms = "http://purl.org/dc/terms/"
                    for field in ("created", "modified"):
                        node = root.find(f"{{{dcterms}}}{field}")
                        if node is not None:
                            node.text = "2026-03-02T09:00:00Z"
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                copied = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
                copied.compress_type = zipfile.ZIP_DEFLATED
                copied.external_attr = info.external_attr
                copied.create_system = info.create_system
                copied.flag_bits = info.flag_bits
                destination.writestr(copied, data)
        shutil.copy2(target, path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for child in args.output_dir.iterdir():
        if child.is_file():
            child.unlink()
        elif child.is_dir():
            shutil.rmtree(child)

    rows = read_sources(args.input_dir)
    schedule = build_schedule(rows)
    build_workbook(rows, schedule, args.output_dir / XLSX_NAME)
    build_program_document(schedule, args.output_dir / DOCX_NAME)
    build_program_page(schedule, args.output_dir / PAGE_NAME)
    build_notifications(rows, schedule, args.output_dir / MBOX_NAME)


if __name__ == "__main__":
    main()
