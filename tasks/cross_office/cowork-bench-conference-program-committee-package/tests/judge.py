#!/usr/bin/env python3
"""ДРВБ-2026 程序委员会包的 v2 连续评分入口；v1 保存在 judge_v1.py。"""
from __future__ import annotations
import argparse,hashlib,importlib.util,json,mailbox,re,shutil,sys
from email import policy
from email.parser import BytesParser
from pathlib import Path
from typing import Any,Iterable
from docx import Document
from openpyxl import load_workbook

TASK_ID="cowork-evo/cowork-bench-conference-program-committee-package";RUBRIC_VERSION="cowork-bench-conference-program-committee-v0.2";JUDGE_VERSION="cowork-bench-conference-program-committee-criterion-partial-v2.0"
XLSX="Submissions.xlsx";DOCX="Conference_Program.docx";PAGE="Program_Page.md";MBOX="Author_Notifications.mbox";REQUIRED=(XLSX,DOCX,PAGE,MBOX)
POSITIVE_WEIGHTS={"R001":5,"R002":12,"R003":12,"R004":12,"R005":8,"R006":10,"R007":8,"R008":12,"R009":5,"R010":8,"R011":4,"R012":4}
PENALTY_WEIGHTS={"R013":-30,"R014":-25,"R015":-30,"R016":-25,"R017":-15}

def _load_v1():
 p=Path(__file__).with_name("judge_v1.py");sys.path.insert(0,str(p.parent)) if str(p.parent) not in sys.path else None;s=importlib.util.spec_from_file_location("conference_judge_v1",p)
 if s is None or s.loader is None:raise RuntimeError("无法载入v1 Judge")
 m=importlib.util.module_from_spec(s);sys.modules[s.name]=m;s.loader.exec_module(m);return m
V1=_load_v1()

def _sha(path:Path):
 d=hashlib.sha256()
 with path.open("rb") as h:
  for chunk in iter(lambda:h.read(1048576),b""):d.update(chunk)
 return d.hexdigest()
def _criterion(key,score,evidence,earned=0,total=1):return {"criterion_id":key,"score":round(max(0,min(1,float(score))),6),"evidence":evidence,"method":"deterministic","raw":{"earned_units":round(earned,6),"total_units":round(total,6)}}
def _ratio(findings:Iterable[tuple[bool,str]]):
 rows=list(findings);earned=sum(bool(ok) for ok,_ in rows);return earned/max(1,len(rows)),[label for ok,label in rows if not ok],earned,len(rows)
def _same(a,b):
 if isinstance(b,float):
  try:return abs(float(a)-b)<=.005
  except (TypeError,ValueError):return False
 if b is None:return a in (None,"")
 return V1.normalize(a)==V1.normalize(b)

def _messages(path):
 box=mailbox.mbox(path,factory=lambda h:BytesParser(policy=policy.default).parse(h),create=False);messages=list(box);box.close();return messages
def _body(message):
 try:return message.get_body(preferencelist=("plain",)).get_content() if message.is_multipart() else message.get_content()
 except Exception:return ""

def _core(artifact:Path):
 if not artifact.is_dir():return False,["输出目录不存在"]
 errors=[]
 for name in REQUIRED:
  path=artifact/name
  if not path.is_file() or path.stat().st_size==0:errors.append(f"核心交付缺失或为空：{name}");continue
  try:
   if name==XLSX:
    book=load_workbook(path,read_only=True,data_only=False)
    if book.sheetnames!=["All_Submissions","Accepted_Program","Control_Summary"]:raise ValueError(f"工作表错误：{book.sheetnames}")
    book.close()
   elif name==DOCX:
    if not Document(path).paragraphs:raise ValueError("Word无段落")
   elif name==PAGE:
    if not path.read_text(encoding="utf-8").strip():raise ValueError("Markdown为空")
   elif not _messages(path):raise ValueError("mbox没有可解析邮件")
  except Exception as exc:errors.append(f"核心交付无法解析：{name}: {type(exc).__name__}: {exc}")
 return not errors,errors
def _fatal(errors,hashes):return {"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":hashes or ["0"*64],"criterion_results":[_criterion(k,0,"；".join(errors),0,0) for k in (*POSITIVE_WEIGHTS,*PENALTY_WEIGHTS)],"reward":0.0,"passed":False,"hurdle_failures":["R001"],"penalties":[]}

def _workbook(path:Path,truth,schedule):
 formulas=load_workbook(path,data_only=False);recalc,log=V1.recalculate_candidate(path)
 if recalc is None:raise RuntimeError(f"LibreOffice重算失败：{log}")
 values=load_workbook(recalc,data_only=True)
 source=[];decision=[];program=[];summary=[]
 all_f= formulas["All_Submissions"];all_v=values["All_Submissions"]
 source.extend([([V1.normalize(c.value) for c in all_f[1]]==V1.ALL_HEADERS,"All_Submissions表头错误"),(all_f.max_row==11 and all_f.max_column==13,"All_Submissions形状错误")])
 for row,wanted in enumerate(truth,start=2):
  expected=[wanted["id"],wanted["author"],wanted["email"],wanted["affiliation"],wanted["title"],wanted["topic"],wanted["abstract"],*wanted["scores"]]
  for col,value in enumerate(expected,start=1):source.append((_same(all_f.cell(row,col).value,value),f"All_Submissions!{all_f.cell(row,col).coordinate}来源字段错误"))
  k=V1.normalize(all_f.cell(row,11).value).upper();l=V1.normalize(all_f.cell(row,12).value).upper();m=V1.normalize(all_f.cell(row,13).value).upper()
  decision.extend([(all_f.cell(row,11).data_type=="f" and f"AVERAGE(H{row}:J{row})" in k,f"K{row}平均分公式错误"),(all_f.cell(row,12).data_type=="f" and f"K{row}>=4" in l,f"L{row}录用阈值公式错误"),(all_f.cell(row,13).data_type=="f" and f"L{row}" in m and f"F{row}" in m,f"M{row}分会场公式错误"),(_same(all_v.cell(row,11).value,wanted["average"]),f"K{row}重算平均分错误"),(_same(all_v.cell(row,12).value,wanted["decision"]),f"L{row}重算决策错误"),(_same(all_v.cell(row,13).value,wanted["section"] or None),f"M{row}重算分会场错误")])
 prog_f=formulas["Accepted_Program"];prog_v=values["Accepted_Program"];program.extend([([V1.normalize(c.value) for c in prog_f[1]]==V1.PROGRAM_HEADERS,"Accepted_Program表头错误"),(prog_f.max_row==8 and prog_f.max_column==9,"Accepted_Program形状错误")])
 schedule_rows=[]
 for row,wanted in enumerate(schedule,start=2):
  got=[prog_v.cell(row,1).value,V1.iso_date(prog_v.cell(row,2).value),V1.hm(prog_v.cell(row,3).value),V1.hm(prog_v.cell(row,4).value),prog_v.cell(row,5).value,prog_v.cell(row,6).value,prog_v.cell(row,7).value,prog_v.cell(row,8).value,prog_v.cell(row,9).value]
  expected=[wanted["section"],wanted["date"].isoformat(),wanted["start"].strftime("%H:%M"),wanted["end"].strftime("%H:%M"),wanted["author"],wanted["email"],wanted["title"],wanted["affiliation"],wanted["id"]]
  for col,(seen,value) in enumerate(zip(got,expected),start=1):program.append((_same(seen,value),f"Accepted_Program!{prog_v.cell(row,col).coordinate}日程事实错误"))
  src=int(wanted["id"])+1;tokens={1:f"M{src}",4:f"C{row}+TIME(0,30,0)",5:f"B{src}",6:f"C{src}",7:f"E{src}",8:f"D{src}",9:f"A{src}"}
  for col,token in tokens.items():program.append((prog_f.cell(row,col).data_type=="f" and token in V1.normalize(prog_f.cell(row,col).value).upper(),f"Accepted_Program!{prog_f.cell(row,col).coordinate}公式引用错误"))
  schedule_rows.append({"section":V1.normalize(got[0]),"date":got[1],"start":got[2],"end":got[3],"author":V1.normalize(got[4]),"email":V1.normalize(got[5]).lower(),"title":V1.normalize(got[6]),"affiliation":V1.normalize(got[7]),"id":got[8]})
 sum_f=formulas["Control_Summary"];sum_v=values["Control_Summary"];labels=[V1.normalize(sum_v.cell(r,1).value) for r in range(2,11)];summary.append((labels==V1.SUMMARY_METRICS,"Control_Summary指标顺序错误"));wanted_values=[10,7,3,3,2,2,"10:00","11:30",0]
 formula_tokens=[("COUNTA","ALL_SUBMISSIONS","A2:A11"),("COUNTIF","ALL_SUBMISSIONS","L2:L11","ПРИНЯТ"),("COUNTIF","ALL_SUBMISSIONS","L2:L11","ОТКЛОНЁН"),("COUNTIF","ACCEPTED_PROGRAM","A2:A8"),("COUNTIF","ACCEPTED_PROGRAM","A2:A8"),("COUNTIF","ACCEPTED_PROGRAM","A2:A8"),("MIN","ACCEPTED_PROGRAM","C2:C8"),("MAX","ACCEPTED_PROGRAM","D2:D8"),("IF","OR","AND","ACCEPTED_PROGRAM")]
 for index,wanted in enumerate(wanted_values,start=2):
  actual=V1.hm(sum_v.cell(index,2).value) if index in (8,9) else sum_v.cell(index,2).value;summary.append((_same(actual,wanted),f"Control_Summary!B{index}重算值错误"));formula=V1.normalize(sum_f.cell(index,2).value).upper();summary.append((sum_f.cell(index,2).data_type=="f" and all(token in formula for token in formula_tokens[index-2]),f"Control_Summary!B{index}公式错误"))
 native=[(sum(len(s.tables) for s in formulas.worksheets)>=3,"Excel原生表少于3"),(all(s.freeze_panes=="A2" for s in formulas.worksheets),"Excel未全部冻结表头"),(all(bool(s.auto_filter.ref) for s in formulas.worksheets),"Excel未全部启用筛选"),(sum(len(s.conditional_formatting) for s in formulas.worksheets)>=2,"Excel条件格式不足"),(sum(len(s.data_validations.dataValidation) for s in formulas.worksheets)>=1,"Excel评分数据验证缺失")]
 formulas.close();values.close();shutil.rmtree(recalc.parents[1],ignore_errors=True);return _ratio(source),_ratio(decision),_ratio(program),_ratio(summary),schedule_rows,native

def _word(path,truth,schedule):
 doc=Document(path);text=V1.word_text(doc);blob=V1.normalize(text).lower();findings=[(len(doc.tables)>=4,"Word原生表少于4张"),(sum(str(p.style.name).startswith("Heading") for p in doc.paragraphs)>=3,"Word标题样式不足")]
 for token in ("Древняя Русь и Волжская Булгария","А.П. Смирнова","17–19 марта 2026","Казанский федеральный университет","Институт международных отношений"):findings.append((token.lower() in blob,f"Word缺会议信息 {token}"))
 for item in schedule:
  for token in (item["author"],item["title"],item["affiliation"],item["start"].strftime("%H:%M"),item["end"].strftime("%H:%M")):findings.append((V1.normalize(token).lower() in blob,f"Word缺录用ID {item['id']} 的 {token}"))
 for item in truth:
  if item["decision"]=="Отклонён":findings.extend([(V1.normalize(item["author"]).lower() not in blob,f"Word泄漏拒稿作者 {item['id']}"),(V1.normalize(item["title"]).lower() not in blob,f"Word泄漏拒稿题目 {item['id']}")])
 for token in ("Рецензент 1","Средняя оценка"):findings.append((token.lower() not in blob,f"Word泄漏内部评分字段 {token}"))
 return (*_ratio(findings),text,doc)

def _page(path,truth,schedule):
 text=path.read_text(encoding="utf-8");lines=text.splitlines();blob=V1.normalize(text).lower();findings=[(bool(lines) and lines[0].strip()=="# Программа конференции ДРВБ-2026","Markdown一级标题错误"),(sum(line.startswith("| ") and not line.startswith("| Время") for line in lines)==7,"Markdown日程行数不是7")]
 for item in schedule:
  for token in (item["author"],item["title"],item["affiliation"],item["section"],f"{item['date'].day} марта 2026",item["start"].strftime("%H:%M"),item["end"].strftime("%H:%M")):findings.append((V1.normalize(token).lower() in blob,f"Markdown缺ID {item['id']} 的 {token}"))
 for item in truth:
  if item["decision"]=="Отклонён":findings.extend([(V1.normalize(item["author"]).lower() not in blob,f"Markdown泄漏拒稿作者 {item['id']}"),(V1.normalize(item["title"]).lower() not in blob,f"Markdown泄漏拒稿题目 {item['id']}")])
 return (*_ratio(findings),text)

def _mail(path,truth,schedule):
 messages=_messages(path);by={}
 for msg in messages:by.setdefault(V1.normalize(msg.get("To","")).lower(),[]).append(msg)
 expected_emails={str(r["email"]).lower() for r in truth};delivery=[(len(messages)==10,"mbox邮件数不是10"),(set(by)==expected_emails,"mbox收件人集合错误")];accepted=[];rejected=[];privacy=[];sched={int(r["id"]):r for r in schedule};records={}
 for row in truth:
  email=str(row["email"]).lower();items=by.get(email,[]);delivery.append((len(items)==1,f"{email}不是恰好一封"))
  if len(items)!=1:continue
  msg=items[0];body=_body(msg);blob=V1.normalize(body).lower();subject=V1.normalize(msg.get("Subject",""));records[row["id"]]={"subject":subject,"body":body,"to":email}
  common=[str(row["author"]).split()[0],row["title"],"ДРВБ-2026","17–19 марта 2026"]
  target=accepted if row["decision"]=="Принят" else rejected
  target.append((subject==("Ваш доклад принят на ДРВБ-2026" if row["decision"]=="Принят" else "Решение по заявке ДРВБ-2026"),f"ID {row['id']}邮件主题错误"))
  for token in common:target.append((V1.normalize(token).lower() in blob,f"ID {row['id']}邮件缺 {token}"))
  if row["decision"]=="Принят":
   item=sched[int(row["id"])]
   for token in (item["section"],item["date"].isoformat(),item["start"].strftime("%H:%M")):accepted.append((V1.normalize(token).lower() in blob,f"ID {row['id']}录用邮件缺 {token}"))
  else:
   rejected.append((bool(re.search(r"не\s+включ[её]н|отклон|не\s+принят",blob,re.I)),f"ID {row['id']}拒稿含义不清"));privacy.append((not re.search(r"реценз|оценк|балл|рейтинг",blob,re.I),f"ID {row['id']}泄漏评分/排名"))
   for other in truth:
    if other["id"]==row["id"]:continue
    for token in (str(other["author"]).split()[0],other["email"],other["title"]):privacy.append((V1.normalize(token).lower() not in blob,f"ID {row['id']}泄漏ID {other['id']} 的 {token}"))
 return _ratio(delivery),_ratio(accepted),_ratio(rejected+privacy),records

def _cross(schedule,word_text,page_text,mail_records):
 w=V1.normalize(word_text).lower();p=V1.normalize(page_text).lower();findings=[]
 for item in schedule:
  mail=mail_records.get(item["id"],{});m=V1.normalize(mail.get("body","")).lower()
  for token in (item["author"],item["title"],item["section"],item["start"].strftime("%H:%M")):findings.append((V1.normalize(token).lower() in w and V1.normalize(token).lower() in p and (V1.normalize(token).lower() in m or token==item["author"]),f"ID {item['id']} 的 {token} 未贯穿Word/页面/邮件"))
  russian_date=f"{item['date'].day} марта 2026";iso_date=item["date"].isoformat();findings.append((russian_date in w and russian_date in p and iso_date in m,f"ID {item['id']} 的日期未按各载体约定保持一致"))
 return _ratio(findings)

def _render_native(artifact,native,doc):
 packages={name:V1.inspect_package(artifact/name) for name in (XLSX,DOCX)};render=V1.render_office(artifact);findings=list(native)+[(len(doc.tables)>=4,"Word原生表少于4张"),(not packages[DOCX]["media"],"Word用图片替代日程"),(all(not info["broken"] and not info["missing_targets"] and not info["macros"] and not info["external"] for info in packages.values()),"Office包损坏、外链或宏异常"),(not render["errors"],f"LibreOffice渲染错误：{render['errors']}"),(2<=render["xlsx_pages"]<=6,"Excel渲染页数异常"),(2<=render["docx_pages"]<=5,"Word渲染页数异常")];return _ratio(findings)

def _identity_penalty(artifact,truth):
 book=load_workbook(artifact/XLSX,data_only=False);ids=[book["All_Submissions"].cell(r,1).value for r in range(2,12)];emails=[V1.normalize(book["All_Submissions"].cell(r,3).value).lower() for r in range(2,12)];book.close();expected_ids={r["id"] for r in truth};expected_emails={str(r["email"]).lower() for r in truth};bad=(set(ids)!=expected_ids or len(ids)!=len(set(ids)) or set(emails)!=expected_emails or len(emails)!=len(set(emails)));return (1.0 if bad else 0.0),(["申请ID或邮箱集合缺失、重复或被替换"] if bad else []),1.0
def _schedule_penalty(schedule_rows,truth):
 rejected={r["id"] for r in truth if r["decision"]=="Отклонён"};ids=[r.get("id") for r in schedule_rows];leaks=[i for i in ids if i in rejected];overlaps=[]
 for a,b in zip(schedule_rows,schedule_rows[1:]):
  if a.get("date")==b.get("date") and a.get("end") and b.get("start") and a["end"]>b["start"]:overlaps.append((a.get("id"),b.get("id")))
 severity=max((1.0 if leaks else 0),(.8 if overlaps else 0));labels=[]
 if leaks:labels.append(f"日程包含拒稿ID：{leaks}")
 if overlaps:labels.append(f"日程时段重叠：{overlaps}")
 return severity,labels,1.0
def _mail_penalty(mail_scores):
 delivery,accepted,rejected=mail_scores;issues=[]
 if delivery<1:issues.append((.8,"邮件错投、缺投、重复或数量错误"))
 if accepted<1:issues.append((.65,"录用通知含义或个人排期错误"))
 if rejected<1:issues.append((1.0,"拒稿通知含义错误或发生隐私泄漏"))
 severity=max((s for s,_ in issues),default=0);return severity,[label for _,label in issues],1.0
def _safe_penalty(artifact):
 delivered={p.name for p in artifact.iterdir()};allowed={*REQUIRED,"__pycache__"};extras=sorted(delivered-allowed);suspicious=V1.suspicious_content(artifact);packages={name:V1.inspect_package(artifact/name) for name in (XLSX,DOCX)};unsafe=any(info["macros"] or info["external"] for info in packages.values());severity=max((1.0 if unsafe else 0),(.8 if suspicious else 0),(.6 if extras else 0));labels=[]
 if unsafe:labels.append("Office含宏或危险外链")
 if suspicious:labels.append(f"绝对路径/Golden/Judge泄漏：{suspicious}")
 if extras:labels.append(f"额外输出：{extras}")
 return severity,labels,1.0
def _bio_penalty(artifact):
 word=V1.parse_word(artifact/DOCX,*V1.load_truth()[:2]);findings=word.get("unsupported",[]);return (1.0 if findings else 0.0),findings,1.0

def evaluate(artifact:Path):
 ok,errors=_core(artifact);hashes=[_sha(artifact/n) for n in REQUIRED if (artifact/n).is_file()]
 if not ok:return _fatal(errors,hashes)
 truth,schedule,source_errors=V1.load_truth()
 if source_errors:return _fatal([f"冻结输入校验失败：{source_errors}"],hashes)
 source,decision,program,summary,schedule_rows,native=_workbook(artifact/XLSX,truth,schedule);word_score,word_fail,word_e,word_t,word_text,doc=_word(artifact/DOCX,truth,schedule);page_score,page_fail,page_e,page_t,page_text=_page(artifact/PAGE,truth,schedule);delivery,accepted,rejected,mail_records=_mail(artifact/MBOX,truth,schedule)
 results=[_criterion("R001",1,"四份核心交付存在、非空且可解析；冻结输入哈希正确",4,4)];scores={"R001":1.0};groups={"R002":source,"R003":decision,"R004":program,"R005":summary,"R006":(word_score,word_fail,word_e,word_t),"R007":(page_score,page_fail,page_e,page_t),"R008":delivery,"R009":accepted,"R010":rejected,"R011":_cross(schedule,word_text,page_text,mail_records),"R012":_render_native(artifact,native,doc)}
 for key in [f"R{i:03d}" for i in range(2,13)]:
  score,fail,earned,total=groups[key];results.append(_criterion(key,score,"全部任务专属原子检查通过" if not fail else "未通过："+"；".join(str(x) for x in fail),earned,total));scores[key]=score
 penalties={"R013":_identity_penalty(artifact,truth),"R014":_schedule_penalty(schedule_rows,truth),"R015":_mail_penalty((delivery[0],accepted[0],rejected[0])),"R016":_safe_penalty(artifact),"R017":_bio_penalty(artifact)}
 for key,(score,labels,total) in penalties.items():results.append(_criterion(key,score,"未触发任务专属严重边界" if not labels else "；".join(str(x) for x in labels),score*total,total));scores[key]=score
 total=sum(POSITIVE_WEIGHTS.values());positive=sum(POSITIVE_WEIGHTS[k]*scores[k] for k in POSITIVE_WEIGHTS);penalty=sum(abs(PENALTY_WEIGHTS[k])*scores[k] for k in PENALTY_WEIGHTS);reward=round(max(0,min(1,(positive-penalty)/total)),6)
 return {"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":hashes,"criterion_results":results,"reward":reward,"passed":reward>=.8,"hurdle_failures":[],"penalties":[{"id":k,"amount":round(abs(PENALTY_WEIGHTS[k])*scores[k]/total,6)} for k in PENALTY_WEIGHTS if scores[k]>0]}

def main():
 p=argparse.ArgumentParser();p.add_argument("--artifact-dir",type=Path,required=True);p.add_argument("--result",type=Path,required=True);p.add_argument("--reward-file",type=Path,required=True);p.add_argument("--scoring-version",choices=("v1","v2"),default="v2");a=p.parse_args()
 try:r=V1.evaluate(a.artifact_dir) if a.scoring_version=="v1" else evaluate(a.artifact_dir)
 except Exception as exc:r=_fatal([f"Judge异常：{type(exc).__name__}: {exc}"],[])
 out=json.dumps(r,ensure_ascii=False,indent=2);a.result.parent.mkdir(parents=True,exist_ok=True);a.reward_file.parent.mkdir(parents=True,exist_ok=True);a.result.write_text(out+"\n",encoding="utf-8");a.reward_file.write_text(f"{float(r['reward']):.6f}\n",encoding="utf-8");print(out)
if __name__=="__main__":main()
