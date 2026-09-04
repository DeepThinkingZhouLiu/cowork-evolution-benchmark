#!/usr/bin/env python3
"""ДРВБ-2026 程序委员会任务的来源重算、跨文件一致性与隐私 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import mailbox
import re
import shutil
import subprocess
import tempfile
import zipfile
from datetime import date, datetime, time, timedelta
from email import policy
from email.parser import BytesParser
from pathlib import Path

from docx import Document
from lxml import etree
from openpyxl import load_workbook


TASK_ID = "cowork-evo/cowork-bench-conference-program-committee-package"
RUBRIC_VERSION = "cowork-bench-conference-program-committee-v0.1"
JUDGE_VERSION = "cowork-bench-conference-program-committee-deterministic-v0.1"
XLSX_NAME = "Submissions.xlsx"
DOCX_NAME = "Conference_Program.docx"
PAGE_NAME = "Program_Page.md"
MBOX_NAME = "Author_Notifications.mbox"
REQUIRED_FILES = [XLSX_NAME, DOCX_NAME, PAGE_NAME, MBOX_NAME]
SOURCE_HASHES = {
    "conference_submissions.json": "e0b94a539a1b0ceb46485215d157b5c27dab69024fd0bc73c3b7cd91f5021775",
    "review_scores.md": "229fb512c5d2221aae62437a5bf5039aeefcc08ceb49ae689c76630c5dd5e6e9",
    "conference_rules.md": "ea8bcfcd6c8e305ce70afa79fd07cb5376ad4dc32177bb90af5cf4900b46fe86",
}
ALL_HEADERS = [
    "ID", "ФИО автора", "Email", "Аффилиация", "Название доклада", "Тематика", "Аннотация",
    "Рецензент 1", "Рецензент 2", "Рецензент 3", "Средняя оценка", "Статус", "Секция",
]
PROGRAM_HEADERS = ["Секция", "День", "Начало", "Окончание", "Докладчик", "Email", "Название доклада", "Аффилиация", "Source ID"]
SUMMARY_METRICS = [
    "Total_Submissions", "Accepted", "Rejected", "Археология", "Источниковедение", "История",
    "Earliest_Start", "Latest_End", "Schedule_Conflicts",
]
SECTION_PLAN = {
    "Археология": (date(2026, 3, 17), time(10, 0)),
    "Источниковедение": (date(2026, 3, 18), time(10, 0)),
    "История": (date(2026, 3, 19), time(10, 0)),
}
SUSPICIOUS_MARKERS = ["solution/golden", "judge.py", "expectations.json", "judge-result", "/root/", "/home/", "c:\\"]

SPECS = {
    "R001": ("完整交付物与冻结输入", 5, "hurdle", 0.0),
    "R002": ("十份申请与三十个评分", 12, "hurdle", 0.3),
    "R003": ("公式化平均分、决策与分会场", 12, "hurdle", 0.35),
    "R004": ("七人三天日程", 12, "hurdle", 0.4),
    "R005": ("控制汇总公式", 8, "hurdle", 0.5),
    "R006": ("Word 正式节目册", 10, "hurdle", 0.5),
    "R007": ("Teamly 页面快照", 8, "hurdle", 0.55),
    "R008": ("十封通知与唯一投递", 12, "hurdle", 0.4),
    "R009": ("录用通知个性化", 5, "hurdle", 0.55),
    "R010": ("拒稿通知与隐私隔离", 8, "hurdle", 0.55),
    "R011": ("四文件事实一致", 4, "hurdle", 0.6),
    "R012": ("原生结构与渲染", 4, "positive", None),
    "R013": ("无来源真值或决策冲突", -30, "penalty", None),
    "R014": ("无排期或跨文件冲突", -25, "penalty", None),
    "R015": ("无错投、误判或隐私泄露", -30, "penalty", None),
    "R016": ("无缺失、损坏、静态化或调试泄漏", -25, "penalty", None),
    "R017": ("无不支持的传记与会议事实", -15, "penalty", None),
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def normalize(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\u00a0", " ").replace("\v", " ")).strip()


def source_path(name: str) -> Path | None:
    task_root = Path(__file__).resolve().parents[1]
    candidates = [
        Path("/app/input_files") / name,
        task_root / "data" / "input_files" / name,
        task_root / "environment" / "input_files" / name,
    ]
    return next((path for path in candidates if path.is_file()), None)


def load_truth() -> tuple[list[dict[str, object]], list[dict[str, object]], dict[str, str]]:
    source_errors: dict[str, str] = {}
    paths: dict[str, Path] = {}
    for name, expected_hash in SOURCE_HASHES.items():
        path = source_path(name)
        if path is None:
            source_errors[name] = "missing"
            continue
        paths[name] = path
        actual = sha256(path)
        if actual != expected_hash:
            source_errors[name] = f"sha256={actual}"
    if source_errors:
        return [], [], source_errors

    payload = json.loads(paths["conference_submissions.json"].read_text(encoding="utf-8"))
    score_map: dict[int, tuple[int, int, int]] = {}
    for line in paths["review_scores.md"].read_text(encoding="utf-8").splitlines():
        match = re.match(r"^\|\s*(\d+)\s*\|.*?\|\s*(\d)\s*\|\s*(\d)\s*\|\s*(\d)\s*\|$", line)
        if match:
            score_map[int(match.group(1))] = tuple(int(match.group(index)) for index in (2, 3, 4))

    rows: list[dict[str, object]] = []
    for response in sorted(payload["responses"], key=lambda item: int(item["id"])):
        identifier = int(response["id"])
        scores = score_map[identifier]
        average = round(sum(scores) / 3, 2)
        accepted = average >= 4.0
        rows.append(
            {
                "id": identifier,
                "author": str(response["ФИО автора"]),
                "email": str(response["Email"]),
                "affiliation": str(response["Аффилиация"]),
                "title": str(response["Название доклада"]),
                "topic": str(response["Тематика"]),
                "abstract": str(response["Аннотация"]),
                "scores": scores,
                "average": average,
                "decision": "Принят" if accepted else "Отклонён",
                "section": str(response["Тематика"]) if accepted else "",
            }
        )

    schedule: list[dict[str, object]] = []
    for section, (event_date, first_start) in SECTION_PLAN.items():
        candidates = [row for row in rows if row["decision"] == "Принят" and row["section"] == section]
        for offset, row in enumerate(sorted(candidates, key=lambda item: int(item["id"]))):
            start_dt = datetime.combine(event_date, first_start) + timedelta(minutes=30 * offset)
            schedule.append({**row, "date": event_date, "start": start_dt.time(), "end": (start_dt + timedelta(minutes=30)).time()})
    schedule.sort(key=lambda item: (item["date"], item["start"], item["id"]))
    return rows, schedule, source_errors


def iso_date(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = normalize(value)
    try:
        return date.fromisoformat(text[:10]).isoformat()
    except ValueError:
        return None


def hm(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.strftime("%H:%M")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    text = normalize(value)
    for pattern in (r"^(\d{1,2}):(\d{2})", r"T(\d{2}):(\d{2})"):
        match = re.search(pattern, text)
        if match:
            return f"{int(match.group(1)):02d}:{match.group(2)}"
    return None


def inspect_package(path: Path) -> dict[str, object]:
    result: dict[str, object] = {"broken": None, "missing_targets": [], "macros": [], "external": [], "media": []}
    if not path.is_file():
        result["broken"] = "missing"
        return result
    try:
        with zipfile.ZipFile(path) as package:
            names = set(package.namelist())
            result["macros"] = sorted(name for name in names if "vbaProject" in name or name.endswith(".bin"))
            result["media"] = sorted(name for name in names if "/media/" in name)
            for rel_name in [name for name in names if name.endswith(".rels")]:
                root = etree.fromstring(package.read(rel_name))
                for rel in root:
                    target = rel.get("Target", "")
                    if rel.get("TargetMode") == "External":
                        result["external"].append(target)
                        continue
                    if target.startswith("/"):
                        resolved = target.lstrip("/")
                    else:
                        base = Path(rel_name).parent
                        if base.name == "_rels":
                            base = base.parent
                        resolved = str((base / target).as_posix())
                        while "/../" in resolved:
                            resolved = re.sub(r"[^/]+/\.\./", "", resolved, count=1)
                    if target and not target.startswith("#") and resolved not in names and not target.startswith("mailto:"):
                        result["missing_targets"].append({"rels": rel_name, "target": target, "resolved": resolved})
    except Exception as exc:
        result["broken"] = f"{type(exc).__name__}: {exc}"
    return result


def recalculate_candidate(path: Path) -> tuple[Path | None, str]:
    temp_root = Path(tempfile.mkdtemp(prefix="conference-judge-recalc-"))
    input_root = temp_root / "input"
    output_root = temp_root / "output"
    profile = temp_root / "profile"
    input_root.mkdir()
    output_root.mkdir()
    profile.mkdir()
    source = input_root / path.name
    shutil.copy2(path, source)
    command = [
        "soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "xlsx",
        "--outdir", str(output_root), str(source),
    ]
    try:
        result = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300)
    except Exception as exc:
        return None, f"{type(exc).__name__}: {exc}"
    converted = output_root / path.name
    if result.returncode != 0 or not converted.is_file():
        return None, result.stdout[-1000:]
    return converted, result.stdout[-1000:]


def parse_workbook(path: Path, truth_rows: list[dict[str, object]], expected_schedule: list[dict[str, object]]) -> dict[str, object]:
    result: dict[str, object] = {
        "readable": False, "sheetnames": [], "source_errors": [], "formula_errors": [], "formula_counts": {},
        "decision_errors": [], "schedule_errors": [], "summary_errors": [], "native": {}, "schedule": [],
    }
    try:
        formulas = load_workbook(path, data_only=False)
    except Exception as exc:
        result["error"] = f"{type(exc).__name__}: {exc}"
        return result
    recalculated, recalc_log = recalculate_candidate(path)
    result["recalc_log"] = recalc_log
    if recalculated is None:
        return result
    try:
        values = load_workbook(recalculated, data_only=True)
    except Exception as exc:
        result["error"] = f"recalc load: {type(exc).__name__}: {exc}"
        return result
    result["readable"] = True
    result["sheetnames"] = formulas.sheetnames
    if formulas.sheetnames != ["All_Submissions", "Accepted_Program", "Control_Summary"]:
        result["source_errors"].append(f"sheetnames={formulas.sheetnames}")
        return result

    all_f = formulas["All_Submissions"]
    all_v = values["All_Submissions"]
    headers = [normalize(cell.value) for cell in all_f[1]]
    if headers != ALL_HEADERS:
        result["source_errors"].append(f"all_headers={headers}")
    if all_f.max_row != 11 or all_f.max_column != 13:
        result["source_errors"].append(f"all_shape={all_f.max_row}x{all_f.max_column}")
    all_formula_count = sum(cell.data_type == "f" for row in all_f.iter_rows() for cell in row)
    result["formula_counts"]["All_Submissions"] = all_formula_count
    for excel_row, expected in enumerate(truth_rows, start=2):
        source_values = [
            expected["id"], expected["author"], expected["email"], expected["affiliation"], expected["title"],
            expected["topic"], expected["abstract"], *expected["scores"],
        ]
        for column, expected_value in enumerate(source_values, start=1):
            got = all_f.cell(excel_row, column).value
            if normalize(got) != normalize(expected_value):
                result["source_errors"].append(f"All_Submissions!{all_f.cell(excel_row, column).coordinate}={got!r}")
        for column in (11, 12, 13):
            cell = all_f.cell(excel_row, column)
            if cell.data_type != "f":
                result["formula_errors"].append(f"{cell.coordinate}:not_formula")
        formula_k = normalize(all_f.cell(excel_row, 11).value).upper()
        formula_l = normalize(all_f.cell(excel_row, 12).value).upper()
        formula_m = normalize(all_f.cell(excel_row, 13).value).upper()
        if f"AVERAGE(H{excel_row}:J{excel_row})" not in formula_k:
            result["formula_errors"].append(f"K{excel_row}:wrong_source")
        if f"K{excel_row}>=4" not in formula_l:
            result["formula_errors"].append(f"L{excel_row}:wrong_threshold")
        if f"L{excel_row}" not in formula_m or f"F{excel_row}" not in formula_m:
            result["formula_errors"].append(f"M{excel_row}:wrong_source")
        cached = [all_v.cell(excel_row, column).value for column in (11, 12, 13)]
        expected_cached = [expected["average"], expected["decision"], expected["section"] or None]
        if abs(float(cached[0] or -999) - float(expected_cached[0])) > 0.005 or normalize(cached[1]) != normalize(expected_cached[1]) or normalize(cached[2]) != normalize(expected_cached[2]):
            result["decision_errors"].append({"id": expected["id"], "got": cached, "expected": expected_cached})

    program_f = formulas["Accepted_Program"]
    program_v = values["Accepted_Program"]
    program_headers = [normalize(cell.value) for cell in program_f[1]]
    if program_headers != PROGRAM_HEADERS:
        result["schedule_errors"].append(f"headers={program_headers}")
    if program_f.max_row != 8 or program_f.max_column != 9:
        result["schedule_errors"].append(f"shape={program_f.max_row}x{program_f.max_column}")
    program_formula_count = sum(cell.data_type == "f" for row in program_f.iter_rows() for cell in row)
    result["formula_counts"]["Accepted_Program"] = program_formula_count
    for excel_row, expected in enumerate(expected_schedule, start=2):
        got = {
            "section": normalize(program_v.cell(excel_row, 1).value),
            "date": iso_date(program_v.cell(excel_row, 2).value),
            "start": hm(program_v.cell(excel_row, 3).value),
            "end": hm(program_v.cell(excel_row, 4).value),
            "author": normalize(program_v.cell(excel_row, 5).value),
            "email": normalize(program_v.cell(excel_row, 6).value),
            "title": normalize(program_v.cell(excel_row, 7).value),
            "affiliation": normalize(program_v.cell(excel_row, 8).value),
            "id": program_v.cell(excel_row, 9).value,
        }
        wanted = {
            "section": normalize(expected["section"]), "date": expected["date"].isoformat(),
            "start": expected["start"].strftime("%H:%M"), "end": expected["end"].strftime("%H:%M"),
            "author": normalize(expected["author"]), "email": normalize(expected["email"]),
            "title": normalize(expected["title"]), "affiliation": normalize(expected["affiliation"]), "id": expected["id"],
        }
        if got != wanted:
            result["schedule_errors"].append({"row": excel_row, "got": got, "expected": wanted})
        result["schedule"].append(got)
        for column in (1, 4, 5, 6, 7, 8, 9):
            if program_f.cell(excel_row, column).data_type != "f":
                result["formula_errors"].append(f"Accepted_Program!{program_f.cell(excel_row, column).coordinate}:not_formula")

    summary_f = formulas["Control_Summary"]
    summary_v = values["Control_Summary"]
    summary_labels = [normalize(summary_v.cell(row, 1).value) for row in range(2, 11)]
    summary_values = [summary_v.cell(row, 2).value for row in range(2, 11)]
    expected_values = [10, 7, 3, 3, 2, 2, "10:00", "11:30", 0]
    result["formula_counts"]["Control_Summary"] = sum(cell.data_type == "f" for row in summary_f.iter_rows() for cell in row)
    if summary_labels != SUMMARY_METRICS:
        result["summary_errors"].append(f"labels={summary_labels}")
    for index, expected in enumerate(expected_values):
        value = summary_values[index]
        got = hm(value) if index in (6, 7) else value
        if got != expected:
            result["summary_errors"].append(f"{SUMMARY_METRICS[index]}={got!r} expected={expected!r}")
        if summary_f.cell(index + 2, 2).data_type != "f":
            result["formula_errors"].append(f"Control_Summary!B{index+2}:not_formula")

    formula_error_cells = []
    for sheet in values.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("#"):
                    formula_error_cells.append(f"{sheet.title}!{cell.coordinate}={cell.value}")
    result["formula_errors"].extend(formula_error_cells)
    result["native"] = {
        "tables": sum(len(sheet.tables) for sheet in formulas.worksheets),
        "freeze": all(sheet.freeze_panes == "A2" for sheet in formulas.worksheets),
        "filters": all(bool(sheet.auto_filter.ref) for sheet in formulas.worksheets),
        "conditional_formats": sum(len(sheet.conditional_formatting) for sheet in formulas.worksheets),
        "data_validations": sum(len(sheet.data_validations.dataValidation) for sheet in formulas.worksheets),
    }
    return result


def word_text(document: Document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def parse_word(path: Path, rows: list[dict[str, object]], schedule: list[dict[str, object]]) -> dict[str, object]:
    result: dict[str, object] = {"readable": False, "errors": [], "unsupported": [], "tables": 0, "headings": []}
    try:
        document = Document(path)
    except Exception as exc:
        result["error"] = f"{type(exc).__name__}: {exc}"
        return result
    result["readable"] = True
    result["tables"] = len(document.tables)
    result["headings"] = [normalize(p.text) for p in document.paragraphs if str(p.style.name).startswith("Heading")]
    text = word_text(document)
    normalized = normalize(text)
    for token in (
        "Древняя Русь и Волжская Булгария", "А.П. Смирнова", "17–19 марта 2026",
        "Казанский федеральный университет", "Институт международных отношений",
    ):
        if token.lower() not in normalized.lower():
            result["errors"].append(f"missing:{token}")
    for expected in schedule:
        for token in (expected["author"], expected["title"], expected["affiliation"]):
            if normalize(token).lower() not in normalized.lower():
                result["errors"].append(f"missing:{expected['id']}:{token}")
    rejected = [row for row in rows if row["decision"] == "Отклонён"]
    for row in rejected:
        if normalize(row["author"]).lower() in normalized.lower() or normalize(row["title"]).lower() in normalized.lower():
            result["errors"].append(f"rejected_leak:{row['id']}")
    schedule_tables = [table for table in document.tables if table.rows and normalize(table.cell(0, 0).text) == "Время"]
    if len(schedule_tables) != 3:
        result["errors"].append(f"schedule_tables={len(schedule_tables)}")
    else:
        for table, section in zip(schedule_tables, SECTION_PLAN, strict=True):
            expected_rows = [item for item in schedule if item["section"] == section]
            if len(table.rows) != len(expected_rows) + 1:
                result["errors"].append(f"{section}:rows={len(table.rows)}")
            for table_row, expected in zip(table.rows[1:], expected_rows, strict=False):
                cells = [normalize(cell.text) for cell in table_row.cells]
                expected_time = f"{expected['start'].strftime('%H:%M')}–{expected['end'].strftime('%H:%M')}"
                if expected_time not in cells[0] or normalize(expected["author"]) not in cells[1] or normalize(expected["affiliation"]) not in cells[1] or normalize(expected["title"]) not in cells[2]:
                    result["errors"].append({"section": section, "cells": cells, "id": expected["id"]})
    for pattern in (r"1899\s*[–-]\s*1974", r"выдающ\w*\s+советск\w*\s+археолог", r"основател\w*\s+школ"):
        if re.search(pattern, normalized, flags=re.IGNORECASE):
            result["unsupported"].append(pattern)
    return result


def parse_page(path: Path, rows: list[dict[str, object]], schedule: list[dict[str, object]]) -> dict[str, object]:
    result: dict[str, object] = {"readable": False, "errors": [], "row_count": 0}
    try:
        text = path.read_text(encoding="utf-8")
    except Exception as exc:
        result["error"] = f"{type(exc).__name__}: {exc}"
        return result
    result["readable"] = True
    lines = text.splitlines()
    if not lines or lines[0].strip() != "# Программа конференции ДРВБ-2026":
        result["errors"].append("wrong_title")
    normalized = normalize(text)
    data_rows = [line for line in lines if line.startswith("| ") and not line.startswith("| Время")]
    result["row_count"] = len(data_rows)
    if len(data_rows) != 7:
        result["errors"].append(f"row_count={len(data_rows)}")
    for expected in schedule:
        expected_tokens = [
            expected["author"], expected["title"], expected["affiliation"], expected["section"],
            str(expected["date"].day) + " марта 2026", expected["start"].strftime("%H:%M"), expected["end"].strftime("%H:%M"),
        ]
        for token in expected_tokens:
            if normalize(token).lower() not in normalized.lower():
                result["errors"].append(f"missing:{expected['id']}:{token}")
    for row in rows:
        if row["decision"] == "Отклонён" and (normalize(row["author"]).lower() in normalized.lower() or normalize(row["title"]).lower() in normalized.lower()):
            result["errors"].append(f"rejected_leak:{row['id']}")
    return result


def parse_mbox(path: Path, rows: list[dict[str, object]], schedule: list[dict[str, object]]) -> dict[str, object]:
    result: dict[str, object] = {
        "readable": False, "count": 0, "delivery_errors": [], "accepted_errors": [], "rejected_errors": [], "privacy_errors": [], "messages": [],
    }
    try:
        box = mailbox.mbox(path, factory=lambda handle: BytesParser(policy=policy.default).parse(handle), create=False)
        messages = list(box)
        box.close()
    except Exception as exc:
        result["error"] = f"{type(exc).__name__}: {exc}"
        return result
    result["readable"] = True
    result["count"] = len(messages)
    by_email: dict[str, list] = {}
    for message in messages:
        recipient = normalize(message.get("To", "")).lower()
        by_email.setdefault(recipient, []).append(message)
        result["messages"].append({"to": recipient, "subject": normalize(message.get("Subject", ""))})
    expected_emails = {str(row["email"]).lower() for row in rows}
    if set(by_email) != expected_emails:
        result["delivery_errors"].append({"got": sorted(by_email), "expected": sorted(expected_emails)})
    for email, items in by_email.items():
        if len(items) != 1:
            result["delivery_errors"].append(f"{email}:count={len(items)}")
    schedule_by_id = {int(item["id"]): item for item in schedule}
    for row in rows:
        items = by_email.get(str(row["email"]).lower(), [])
        if len(items) != 1:
            continue
        message = items[0]
        subject = normalize(message.get("Subject", ""))
        try:
            body = message.get_body(preferencelist=("plain",)).get_content() if message.is_multipart() else message.get_content()
        except Exception:
            body = ""
        body_norm = normalize(body)
        surname = str(row["author"]).split()[0]
        common_tokens = [surname, row["title"], "ДРВБ-2026", "17–19 марта 2026"]
        if row["decision"] == "Принят":
            expected = schedule_by_id[int(row["id"])]
            if subject != "Ваш доклад принят на ДРВБ-2026":
                result["accepted_errors"].append(f"{row['id']}:subject={subject}")
            for token in [*common_tokens, expected["section"], expected["date"].isoformat(), expected["start"].strftime("%H:%M")]:
                if normalize(token).lower() not in body_norm.lower():
                    result["accepted_errors"].append(f"{row['id']}:missing:{token}")
        else:
            if subject != "Решение по заявке ДРВБ-2026":
                result["rejected_errors"].append(f"{row['id']}:subject={subject}")
            for token in common_tokens:
                if normalize(token).lower() not in body_norm.lower():
                    result["rejected_errors"].append(f"{row['id']}:missing:{token}")
            if not re.search(r"не\s+включ[её]н|отклон|не\s+принят", body_norm, flags=re.IGNORECASE):
                result["rejected_errors"].append(f"{row['id']}:no_rejection_meaning")
            if re.search(r"реценз|оценк|балл|рейтинг", body_norm, flags=re.IGNORECASE):
                result["privacy_errors"].append(f"{row['id']}:score_language")
            for other in rows:
                if other["id"] == row["id"]:
                    continue
                for token in (str(other["author"]).split()[0], other["email"], other["title"]):
                    if normalize(token).lower() in body_norm.lower():
                        result["privacy_errors"].append(f"{row['id']}:leaks:{other['id']}:{token}")
    if len(messages) != 10:
        result["delivery_errors"].append(f"message_count={len(messages)}")
    return result


def render_office(artifact_dir: Path) -> dict[str, object]:
    result: dict[str, object] = {"xlsx_pages": 0, "docx_pages": 0, "errors": [], "logs": []}
    with tempfile.TemporaryDirectory(prefix="conference-judge-render-") as tmp:
        root = Path(tmp)
        output = root / "output"
        profile = root / "profile"
        output.mkdir()
        profile.mkdir()
        for name in (XLSX_NAME, DOCX_NAME):
            source = artifact_dir / name
            if not source.is_file():
                result["errors"].append(f"missing:{name}")
                continue
            command = [
                "soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf",
                "--outdir", str(output), str(source),
            ]
            try:
                proc = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300)
            except Exception as exc:
                result["errors"].append(f"{name}:{type(exc).__name__}:{exc}")
                continue
            result["logs"].append(proc.stdout[-1000:])
            pdf = output / f"{Path(name).stem}.pdf"
            if proc.returncode != 0 or not pdf.is_file():
                result["errors"].append(f"{name}:conversion_failed")
                continue
            info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=60)
            match = re.search(r"^Pages:\s+(\d+)", info.stdout, flags=re.MULTILINE)
            pages = int(match.group(1)) if match else 0
            text_result = subprocess.run(["pdftotext", str(pdf), "-"], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=60)
            rendered_text = text_result.stdout
            if name == XLSX_NAME:
                result["xlsx_pages"] = pages
                # 工作表名不是正文，不应强制打印到页面上；检查三页的真实内容锨点。
                for token in ("ФИО автора", "Хабибуллин", "Total_Submissions"):
                    if token not in rendered_text:
                        result["errors"].append(f"xlsx_render_missing:{token}")
            else:
                result["docx_pages"] = pages
                for token in ("Древняя Русь", "17 марта 2026", "19 марта 2026", "Площадка"):
                    if token not in rendered_text:
                        result["errors"].append(f"docx_render_missing:{token}")
            if re.search(r"#(?:NAME|VALUE|REF|DIV/0|N/A)", rendered_text, flags=re.IGNORECASE):
                result["errors"].append(f"{name}:formula_error_visible")
    if not 2 <= result["xlsx_pages"] <= 6:
        result["errors"].append(f"xlsx_pages={result['xlsx_pages']}")
    if not 2 <= result["docx_pages"] <= 5:
        result["errors"].append(f"docx_pages={result['docx_pages']}")
    return result


def suspicious_content(artifact_dir: Path) -> list[str]:
    findings: list[str] = []
    for path in artifact_dir.iterdir() if artifact_dir.is_dir() else []:
        if not path.is_file():
            continue
        try:
            if path.suffix.lower() in {".md", ".txt", ".mbox", ".eml", ".ics"}:
                text = path.read_text(encoding="utf-8", errors="ignore").lower()
            elif path.suffix.lower() in {".xlsx", ".docx"}:
                with zipfile.ZipFile(path) as package:
                    text = " ".join(package.read(name).decode("utf-8", errors="ignore") for name in package.namelist() if name.endswith(".xml")).lower()
            else:
                text = ""
        except Exception:
            continue
        for marker in SUSPICIOUS_MARKERS:
            if marker in text:
                findings.append(f"{path.name}:{marker}")
    return sorted(set(findings))


def criterion(identifier: str, score: float, evidence: object) -> dict[str, object]:
    description, weight, kind, cap = SPECS[identifier]
    raw: dict[str, object] = {
        "description": description, "weight": weight, "type": kind, "passed": score >= 0.999,
    }
    if cap is not None:
        raw["score_cap_on_failure"] = cap
    return {"criterion_id": identifier, "score": round(float(score), 6), "evidence": evidence, "method": "deterministic", "raw": raw}


def evaluate(artifact_dir: Path) -> dict[str, object]:
    rows, schedule, source_errors = load_truth()
    delivered = sorted(path.name for path in artifact_dir.iterdir() if path.is_file()) if artifact_dir.is_dir() else []
    exact_files = delivered == sorted(REQUIRED_FILES)
    packages = {
        name: inspect_package(artifact_dir / name)
        for name in (XLSX_NAME, DOCX_NAME)
    }
    package_ok = all(not info["broken"] and not info["missing_targets"] and not info["macros"] and not info["external"] for info in packages.values())

    workbook = parse_workbook(artifact_dir / XLSX_NAME, rows, schedule) if (artifact_dir / XLSX_NAME).is_file() and not source_errors else {"readable": False, "source_errors": ["source_unavailable"], "formula_errors": [], "decision_errors": [], "schedule_errors": [], "summary_errors": [], "native": {}, "formula_counts": {}, "schedule": []}
    word = parse_word(artifact_dir / DOCX_NAME, rows, schedule) if (artifact_dir / DOCX_NAME).is_file() and not source_errors else {"readable": False, "errors": ["unavailable"], "unsupported": [], "tables": 0, "headings": []}
    page = parse_page(artifact_dir / PAGE_NAME, rows, schedule) if (artifact_dir / PAGE_NAME).is_file() and not source_errors else {"readable": False, "errors": ["unavailable"], "row_count": 0}
    messages = parse_mbox(artifact_dir / MBOX_NAME, rows, schedule) if (artifact_dir / MBOX_NAME).is_file() and not source_errors else {"readable": False, "count": 0, "delivery_errors": ["unavailable"], "accepted_errors": [], "rejected_errors": [], "privacy_errors": [], "messages": []}
    render = render_office(artifact_dir) if exact_files else {"xlsx_pages": 0, "docx_pages": 0, "errors": ["incomplete"], "logs": []}
    suspicious = suspicious_content(artifact_dir)

    native = workbook.get("native", {})
    native_ok = (
        native.get("tables", 0) >= 3 and native.get("freeze") is True and native.get("filters") is True
        and native.get("conditional_formats", 0) >= 2 and native.get("data_validations", 0) >= 1
        and word.get("tables", 0) >= 4 and not packages[DOCX_NAME].get("media")
    )
    source_exact = not source_errors and workbook.get("readable") and not workbook.get("source_errors")
    decision_exact = source_exact and not workbook.get("decision_errors") and not workbook.get("formula_errors") and workbook.get("formula_counts", {}).get("All_Submissions", 0) >= 30
    schedule_exact = workbook.get("readable") and not workbook.get("schedule_errors") and workbook.get("formula_counts", {}).get("Accepted_Program", 0) >= 49
    summary_exact = workbook.get("readable") and not workbook.get("summary_errors") and workbook.get("formula_counts", {}).get("Control_Summary", 0) >= 9
    word_ok = word.get("readable") and not word.get("errors") and not word.get("unsupported")
    page_ok = page.get("readable") and not page.get("errors") and page.get("row_count") == 7
    delivery_ok = messages.get("readable") and messages.get("count") == 10 and not messages.get("delivery_errors")
    accepted_mail_ok = delivery_ok and not messages.get("accepted_errors")
    rejected_mail_ok = delivery_ok and not messages.get("rejected_errors") and not messages.get("privacy_errors")
    cross_ok = schedule_exact and word_ok and page_ok and accepted_mail_ok and rejected_mail_ok
    render_ok = not render.get("errors")
    safe_ok = exact_files and package_ok and not suspicious
    unsupported_ok = not word.get("unsupported")

    results = [
        criterion("R001", 1.0 if exact_files and not source_errors and package_ok else 0.0, {"source_errors": source_errors, "delivered": delivered, "packages": packages}),
        criterion("R002", 1.0 if source_exact else 0.0, {"source_errors": workbook.get("source_errors"), "shape": workbook.get("sheetnames")}),
        criterion("R003", 1.0 if decision_exact else 0.0, {"formula_counts": workbook.get("formula_counts"), "formula_errors": workbook.get("formula_errors"), "decision_errors": workbook.get("decision_errors")}),
        criterion("R004", 1.0 if schedule_exact else 0.0, {"schedule_errors": workbook.get("schedule_errors"), "schedule": workbook.get("schedule")}),
        criterion("R005", 1.0 if summary_exact else 0.0, {"summary_errors": workbook.get("summary_errors"), "formula_counts": workbook.get("formula_counts")}),
        criterion("R006", 1.0 if word_ok else 0.0, word),
        criterion("R007", 1.0 if page_ok else 0.0, page),
        criterion("R008", 1.0 if delivery_ok else 0.0, {"count": messages.get("count"), "delivery_errors": messages.get("delivery_errors"), "messages": messages.get("messages")}),
        criterion("R009", 1.0 if accepted_mail_ok else 0.0, messages.get("accepted_errors")),
        criterion("R010", 1.0 if rejected_mail_ok else 0.0, {"rejected_errors": messages.get("rejected_errors"), "privacy_errors": messages.get("privacy_errors")}),
        criterion("R011", 1.0 if cross_ok else 0.0, {"schedule": schedule_exact, "word": word_ok, "page": page_ok, "accepted_mail": accepted_mail_ok, "rejected_mail": rejected_mail_ok}),
        criterion("R012", 1.0 if native_ok and render_ok else 0.0, {"native": native, "word_tables": word.get("tables"), "render": render}),
        criterion("R013", 1.0 if source_exact and decision_exact else 0.0, {"source_exact": source_exact, "decision_exact": decision_exact}),
        criterion("R014", 1.0 if schedule_exact and cross_ok else 0.0, {"schedule_exact": schedule_exact, "cross_ok": cross_ok}),
        criterion("R015", 1.0 if delivery_ok and accepted_mail_ok and rejected_mail_ok else 0.0, {"delivery": delivery_ok, "accepted": accepted_mail_ok, "rejected": rejected_mail_ok, "privacy": messages.get("privacy_errors")}),
        criterion("R016", 1.0 if safe_ok and native_ok else 0.0, {"safe": safe_ok, "native": native_ok, "suspicious": suspicious}),
        criterion("R017", 1.0 if unsupported_ok else 0.0, word.get("unsupported")),
    ]

    positive_points = sum(SPECS[item["criterion_id"]][1] * float(item["score"]) for item in results if SPECS[item["criterion_id"]][1] > 0)
    penalty_points = sum(SPECS[item["criterion_id"]][1] for item in results if SPECS[item["criterion_id"]][1] < 0 and float(item["score"]) < 0.999)
    reward = max(0.0, min(1.0, (positive_points + penalty_points) / 100.0))
    hurdle_failures: list[str] = []
    for item in results:
        identifier = item["criterion_id"]
        _, _, kind, cap = SPECS[identifier]
        if kind == "hurdle" and float(item["score"]) < 0.999:
            hurdle_failures.append(identifier)
            reward = min(reward, float(cap))
    reward = round(reward, 6)
    artifact_hashes = [sha256(artifact_dir / name) for name in REQUIRED_FILES if (artifact_dir / name).is_file()]
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": artifact_hashes,
        "criterion_results": results,
        "reward": reward,
        "passed": reward >= 0.8 and not hurdle_failures,
        "hurdle_failures": hurdle_failures,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    result = evaluate(args.artifact_dir)
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{result['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
