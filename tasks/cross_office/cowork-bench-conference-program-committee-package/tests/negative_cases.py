#!/usr/bin/env python3
"""程序委员会任务的任务专属 v2 致命/局部反例。"""
from __future__ import annotations
import importlib.util,json,mailbox,shutil,sys,tempfile
from email import policy
from email.parser import BytesParser
from pathlib import Path
from docx import Document
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1];GOLDEN=ROOT/"solution"/"Golden";REQUIRED=("Submissions.xlsx","Conference_Program.docx","Program_Page.md","Author_Notifications.mbox")
def _load():
 p=Path(__file__).with_name("judge.py");s=importlib.util.spec_from_file_location("conference_v2_negative_judge",p)
 if s is None or s.loader is None:raise RuntimeError("无法载入Judge")
 m=importlib.util.module_from_spec(s);sys.modules[s.name]=m;s.loader.exec_module(m);return m
JUDGE=_load()
def copy(root):
 root.mkdir(parents=True,exist_ok=True)
 for n in REQUIRED:shutil.copy2(GOLDEN/n,root/n)
def excel(root,sheet,cell,value):
 p=root/"Submissions.xlsx";b=load_workbook(p,data_only=False);b[sheet][cell]=value;b.save(p);b.close()
def score(root):excel(root,"All_Submissions","H2",1)
def title(root):excel(root,"All_Submissions","E2","Неверное название")
def avg_formula(root):excel(root,"All_Submissions","K2",4.67)
def decision_formula(root):excel(root,"All_Submissions","L2",'=IF(K2>=5,"Принят","Отклонён")')
def section_formula(root):excel(root,"All_Submissions","M2","Археология")
def schedule_start(root):excel(root,"Accepted_Program","C2","10:30")
def schedule_date(root):excel(root,"Accepted_Program","B2","2026-03-18")
def schedule_author_formula(root):excel(root,"Accepted_Program","E2","Неизвестный")
def summary_total(root):excel(root,"Control_Summary","B2",10)
def summary_conflict(root):excel(root,"Control_Summary","B10","=1")
def replace_doc(root,old,new):
 p=root/"Conference_Program.docx";d=Document(p)
 for paragraph in [*d.paragraphs,*[p for t in d.tables for r in t.rows for c in r.cells for p in c.paragraphs]]:
  if old in paragraph.text:paragraph.text=paragraph.text.replace(old,new)
 d.save(p)
def word_missing(root):replace_doc(root,"Хабибуллин М.Х.","Неизвестный докладчик")
def word_rejected(root):
 p=root/"Conference_Program.docx";d=Document(p);d.add_paragraph("Краснов В.П. — Тайна копья Святого Олега");d.save(p)
def word_biography(root):
 p=root/"Conference_Program.docx";d=Document(p);d.add_paragraph("А.П. Смирнов (1899–1974) — выдающийся советский археолог и основатель школы.");d.save(p)
def page_title(root):
 p=root/"Program_Page.md";text=p.read_text(encoding="utf-8");p.write_text(text.replace("Археология средневекового Болгара: новые раскопки 2024–2025 гг.","Неверное название доклада",1),encoding="utf-8")
def page_rejected(root):
 p=root/"Program_Page.md";p.write_text(p.read_text(encoding="utf-8")+"\nКраснов В.П. — Тайна копья Святого Олега\n",encoding="utf-8")
def read_mbox(path):
 b=mailbox.mbox(path,factory=lambda h:BytesParser(policy=policy.default).parse(h),create=False);m=list(b);b.close();return m
def write_mbox(path,messages):
 t=path.with_suffix(".new");b=mailbox.mbox(t,create=True)
 for m in messages:b.add(m)
 b.flush();b.close();t.replace(path)
def body_change(message,old=None,new=None,append=None):
 body=message.get_content();body=body.replace(old,new) if old is not None else body;body=body.rstrip()+"\n\n"+append+"\n" if append else body;message.set_content(body,subtype="plain",charset="utf-8")
def mail_missing(root):
 p=root/"Author_Notifications.mbox";write_mbox(p,read_mbox(p)[:-1])
def mail_time(root):
 p=root/"Author_Notifications.mbox";m=read_mbox(p);body_change(m[0],"10:00","12:00");write_mbox(p,m)
def mail_privacy(root):
 p=root/"Author_Notifications.mbox";m=read_mbox(p);target=next(x for x in m if str(x.get("To","")).strip().lower()=="krasnov@example.ru");body_change(target,append="Оценка рецензентов: 2.67 балла.");write_mbox(p,m)
def extra(root):(root/"debug.txt").write_text("temporary\n",encoding="utf-8")
PARTIAL={"one_review_score_wrong":score,"one_submission_title_wrong":title,"average_formula_flattened":avg_formula,"decision_threshold_formula_wrong":decision_formula,"section_formula_flattened":section_formula,"one_schedule_start_wrong":schedule_start,"one_schedule_date_wrong":schedule_date,"one_program_author_formula_flattened":schedule_author_formula,"summary_total_formula_flattened":summary_total,"summary_conflict_formula_wrong":summary_conflict,"word_one_accepted_author_missing":word_missing,"word_rejected_author_leak":word_rejected,"word_unsupported_biography":word_biography,"page_one_accepted_title_wrong":page_title,"page_rejected_author_leak":page_rejected,"mbox_one_message_missing":mail_missing,"accepted_mail_time_wrong":mail_time,"rejected_mail_score_leak":mail_privacy,"unexpected_debug_file":extra}
def main():
 assert JUDGE.evaluate(GOLDEN)["reward"]==1
 JUDGE.V1.render_office=lambda _artifact:{"xlsx_pages":3,"docx_pages":2,"errors":[],"logs":["negative regression uses validated Golden render baseline"]}
 records=[]
 with tempfile.TemporaryDirectory(prefix="conference-v2-neg-") as raw:
  base=Path(raw)
  for n in REQUIRED:
   for mode in ("missing","empty","corrupt"):
    case=base/f"fatal-{mode}-{n}";copy(case);target=case/n
    if mode=="missing":target.unlink()
    elif mode=="empty":target.write_bytes(b"")
    else:target.write_bytes(b"\xff\xfe\x00")
    reward=JUDGE.evaluate(case)["reward"];assert reward==0,(mode,n,reward);records.append({"case":f"{mode}:{n}","class":"fatal","reward":reward})
  case=base/"fatal-empty-output";case.mkdir();reward=JUDGE.evaluate(case)["reward"];assert reward==0;records.append({"case":"empty-output","class":"fatal","reward":reward})
  case=base/"fatal-input-only";shutil.copytree(ROOT/"data"/"input_files",case);reward=JUDGE.evaluate(case)["reward"];assert reward==0;records.append({"case":"input-only","class":"fatal","reward":reward})
  for name,mutate in PARTIAL.items():
   case=base/f"partial-{name}";copy(case);mutate(case);reward=JUDGE.evaluate(case)["reward"];assert 0<reward<1,(name,reward);records.append({"case":name,"class":"partial","reward":reward})
 fatal=[r["reward"] for r in records if r["class"]=="fatal"];partial=[r["reward"] for r in records if r["class"]=="partial"];print(json.dumps({"fatal_count":len(fatal),"fatal_unique_rewards":sorted(set(fatal)),"partial_count":len(partial),"partial_min":min(partial),"partial_max":max(partial),"cases":records},ensure_ascii=False,indent=2))
if __name__=="__main__":main()
