#!/usr/bin/env python3
"""逐个构造本题的坏交付，确认 Judge 能拒绝事实、排期、隐私与文件结构错误。"""

from __future__ import annotations

import argparse
import mailbox
import shutil
import tempfile
from email import policy
from email.parser import BytesParser
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from judge import DOCX_NAME, MBOX_NAME, PAGE_NAME, REQUIRED_FILES, XLSX_NAME, evaluate


CASES = (
    "missing-mbox",
    "formulas-flattened",
    "wrong-source-score",
    "wrong-schedule",
    "word-missing-author",
    "page-rejected-leak",
    "mbox-missing-message",
    "accepted-time-wrong",
    "rejected-score-leak",
    "cross-file-title-conflict",
    "corrupt-docx",
    "extra-debug-file",
)


def copy_golden(golden: Path, case_dir: Path) -> None:
    case_dir.mkdir(parents=True, exist_ok=True)
    for name in REQUIRED_FILES:
        shutil.copy2(golden / name, case_dir / name)


def replace_docx_text(path: Path, old: str, new: str) -> None:
    document = Document(path)
    containers = [document.paragraphs]
    containers.extend(cell.paragraphs for table in document.tables for row in table.rows for cell in row.cells)
    for paragraphs in containers:
        for paragraph in paragraphs:
            if old not in paragraph.text:
                continue
            # 表格单元格中的文字可被分成多个 run，直接重建该段落避免漏改。
            paragraph.text = paragraph.text.replace(old, new)
    document.save(path)


def read_messages(path: Path) -> list:
    box = mailbox.mbox(path, factory=lambda handle: BytesParser(policy=policy.default).parse(handle), create=False)
    messages = list(box)
    box.close()
    return messages


def write_messages(path: Path, messages: list) -> None:
    temporary = path.with_name(path.name + ".new")
    if temporary.exists():
        temporary.unlink()
    box = mailbox.mbox(temporary, create=True)
    for message in messages:
        box.add(message)
    box.flush()
    box.close()
    temporary.replace(path)


def mutate_message_body(message, old: str | None = None, new: str | None = None, append: str | None = None) -> None:
    body = message.get_content()
    if old is not None:
        body = body.replace(old, new or "")
    if append is not None:
        body = body.rstrip() + "\n\n" + append + "\n"
    message.set_content(body, subtype="plain", charset="utf-8")


def mutate(case: str, case_dir: Path) -> None:
    if case == "missing-mbox":
        (case_dir / MBOX_NAME).unlink()
    elif case == "formulas-flattened":
        path = case_dir / XLSX_NAME
        formulas = load_workbook(path, data_only=False)
        values = load_workbook(path, data_only=True)
        for formula_sheet, value_sheet in zip(formulas.worksheets, values.worksheets, strict=True):
            for row in formula_sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        cell.value = value_sheet[cell.coordinate].value
        formulas.save(path)
    elif case == "wrong-source-score":
        path = case_dir / XLSX_NAME
        workbook = load_workbook(path, data_only=False)
        workbook["All_Submissions"]["H2"] = 1
        workbook.save(path)
    elif case == "wrong-schedule":
        path = case_dir / XLSX_NAME
        workbook = load_workbook(path, data_only=False)
        # 第一位录用者原定 10:00，改成与后一位冲突的 10:30。
        workbook["Accepted_Program"]["C2"] = "10:30"
        workbook.save(path)
    elif case == "word-missing-author":
        replace_docx_text(case_dir / DOCX_NAME, "Хабибуллин М.Х.", "Неизвестный докладчик")
    elif case == "page-rejected-leak":
        path = case_dir / PAGE_NAME
        with path.open("a", encoding="utf-8") as handle:
            handle.write("\nОтклонённая заявка: Краснов В.П. — «Тайна копья Святого Олега».\n")
    elif case == "mbox-missing-message":
        path = case_dir / MBOX_NAME
        write_messages(path, read_messages(path)[:-1])
    elif case == "accepted-time-wrong":
        path = case_dir / MBOX_NAME
        messages = read_messages(path)
        mutate_message_body(messages[0], old="10:00", new="12:00")
        write_messages(path, messages)
    elif case == "rejected-score-leak":
        path = case_dir / MBOX_NAME
        messages = read_messages(path)
        rejected = next(message for message in messages if str(message.get("To", "")).strip().lower() == "krasnov@example.ru")
        mutate_message_body(rejected, append="Оценка рецензентов: 2.67 балла.")
        write_messages(path, messages)
    elif case == "cross-file-title-conflict":
        path = case_dir / PAGE_NAME
        text = path.read_text(encoding="utf-8")
        text = text.replace("Археология средневекового Болгара: новые раскопки 2024–2025 гг.", "Неверное название доклада", 1)
        path.write_text(text, encoding="utf-8")
    elif case == "corrupt-docx":
        (case_dir / DOCX_NAME).write_bytes(b"not-a-docx")
    elif case == "extra-debug-file":
        (case_dir / "debug.txt").write_text("temporary judge notes", encoding="utf-8")
    else:
        raise ValueError(case)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--case", choices=CASES, action="append", required=True)
    args = parser.parse_args()
    failures: list[dict[str, object]] = []
    for case in args.case:
        with tempfile.TemporaryDirectory(prefix=f"conference-negative-{case}-") as tmp:
            case_dir = Path(tmp)
            copy_golden(args.golden, case_dir)
            mutate(case, case_dir)
            result = evaluate(case_dir)
            print(f"{case}\treward={result['reward']:.6f}\tpassed={result['passed']}\thurdles={','.join(result['hurdle_failures'])}")
            if result["passed"] or float(result["reward"]) >= 0.8:
                failures.append({"case": case, "result": result})
    if failures:
        raise SystemExit(f"负例错误通过：{[item['case'] for item in failures]}")


if __name__ == "__main__":
    main()
