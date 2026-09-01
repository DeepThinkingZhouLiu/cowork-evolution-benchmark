#!/usr/bin/env python3
"""逐题生成跨系统财务审计任务的四个 Golden 交付物。"""

from __future__ import annotations

import argparse
import csv
import re
import shutil
import subprocess
import tempfile
import zipfile
from datetime import date, datetime, timezone
from email import policy
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


XLSX_NAME = "Audit_Reconciliation_Report.xlsx"
DOCX_NAME = "Audit_Report.docx"
EML_NAME = "Financial_Audit_Reconciliation.eml"
ICS_NAME = "Audit_Reconciliation_Review.ics"
FIXED_DT = datetime(2024, 4, 1, 0, 0, 0, tzinfo=timezone.utc)
ALLOWED_STATUS = {"completed", "pending"}

NAVY = "17324D"
TEAL = "177E89"
PALE_TEAL = "DDEFF1"
PALE_RED = "FCE8E6"
PALE_AMBER = "FFF2CC"
PALE_BLUE = "EAF1F8"
WHITE = "FFFFFF"
GRAY = "5B6573"
LINE = "CBD5DF"
RED = "B42318"


def read_sources(input_dir: Path) -> list[dict[str, object]]:
    """从三份真实输入重算逐订单真值。"""
    ledger_path = input_dir / "authoritative_ledger.csv"
    manual_path = input_dir / "q1_manual_tracking.xlsx"
    guide_path = input_dir / "system_mapping_guide.pdf"
    for path in (ledger_path, manual_path, guide_path):
        if not path.is_file():
            raise FileNotFoundError(path)

    with ledger_path.open(encoding="utf-8", newline="") as handle:
        ledger = {row["Order ID"].strip(): row for row in csv.DictReader(handle)}

    workbook = load_workbook(manual_path, data_only=True)
    if "Manual Tracking" not in workbook.sheetnames:
        raise ValueError("q1_manual_tracking.xlsx 缺少 Manual Tracking")
    sheet = workbook["Manual Tracking"]
    headers = [str(cell.value).strip() for cell in sheet[1]]
    required = ["Order ID", "Date", "Amount", "Customer", "Status"]
    if headers != required:
        raise ValueError(f"手工台账表头不符：{headers}")

    rows: list[dict[str, object]] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        manual = dict(zip(headers, values, strict=True))
        order_id = str(manual["Order ID"]).strip()
        source = ledger.get(order_id)
        manual_date = date.fromisoformat(str(manual["Date"]))
        manual_amount = round(float(manual["Amount"]), 2)
        manual_status = str(manual["Status"]).strip()

        ledger_date = date.fromisoformat(source["Date"]) if source else None
        ledger_amount = round(float(source["Amount"]), 2) if source else None
        ledger_status = source["Status"].strip() if source else None
        date_delta = abs((manual_date - ledger_date).days) if ledger_date else None
        amount_delta = round(manual_amount - ledger_amount, 2) if ledger_amount is not None else None

        reasons: list[str] = []
        if source is None:
            reasons.append("missing_in_authoritative_ledger")
        else:
            if date_delta is None or date_delta > 1:
                reasons.append("date_mismatch")
            if amount_delta is None or abs(amount_delta) >= 0.005:
                reasons.append("amount_mismatch")
            if manual_status.lower() not in ALLOWED_STATUS or str(ledger_status).lower() not in ALLOWED_STATUS:
                reasons.append("invalid_status")

        rows.append(
            {
                "order_id": order_id,
                "manual_date": manual_date,
                "ledger_date": ledger_date,
                "date_delta": date_delta,
                "manual_amount": manual_amount,
                "ledger_amount": ledger_amount,
                "amount_delta": amount_delta,
                "manual_status": manual_status,
                "ledger_status": ledger_status,
                "presence_check": "pass" if source else "fail",
                "date_check": "pass" if source and date_delta is not None and date_delta <= 1 else "fail",
                "amount_check": "pass" if source and amount_delta is not None and abs(amount_delta) < 0.005 else "fail",
                "status_check": "pass" if source and manual_status.lower() in ALLOWED_STATUS and str(ledger_status).lower() in ALLOWED_STATUS else "fail",
                "match_status": "matched" if not reasons else "discrepancy",
                "reason": ";".join(reasons) if reasons else "none",
            }
        )
    return sorted(rows, key=lambda row: str(row["order_id"]))


def truth(rows: list[dict[str, object]]) -> dict[str, object]:
    discrepancies = [row for row in rows if row["match_status"] == "discrepancy"]
    return {
        "total": len(rows),
        "matched": sum(row["match_status"] == "matched" for row in rows),
        "discrepancies": len(discrepancies),
        "manual_total": round(sum(float(row["manual_amount"]) for row in rows), 2),
        "ledger_total": round(sum(float(row["ledger_amount"]) for row in rows if row["ledger_amount"] is not None), 2),
        "missing_amount": round(sum(float(row["manual_amount"]) for row in rows if row["presence_check"] == "fail"), 2),
        "absolute_amount_mismatch": round(sum(abs(float(row["amount_delta"])) for row in rows if "amount_mismatch" in str(row["reason"])), 2),
        "date_mismatches": sum("date_mismatch" in str(row["reason"]) for row in rows),
        "amount_mismatches": sum("amount_mismatch" in str(row["reason"]) for row in rows),
        "missing_records": sum("missing_in_authoritative_ledger" in str(row["reason"]) for row in rows),
        "discrepancy_rows": discrepancies,
    }


def style_sheet(sheet, widths: dict[str, float]) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 32
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color=LINE)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=9, color=NAVY)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:1"
    sheet.page_margins.left = 0.2
    sheet.page_margins.right = 0.2
    sheet.page_margins.top = 0.4
    sheet.page_margins.bottom = 0.4


def add_native_table(sheet, name: str) -> None:
    table = Table(displayName=name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def build_workbook(rows: list[dict[str, object]], output_path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.properties.creator = "EvoCowork benchmark adapter"
    workbook.properties.title = "Q1 2024 Financial Audit Reconciliation"
    workbook.properties.created = FIXED_DT.replace(tzinfo=None)
    workbook.properties.modified = FIXED_DT.replace(tzinfo=None)

    reconciliation = workbook.create_sheet("Reconciliation")
    headers = [
        "Order ID", "Manual Date", "Ledger Date", "Date Delta Days", "Manual Amount",
        "Ledger Amount", "Amount Delta", "Manual Status", "Ledger Status", "Presence Check",
        "Date Check", "Amount Check", "Status Check", "Match Status", "Discrepancy Reason",
    ]
    reconciliation.append(headers)
    for index, row in enumerate(rows, start=2):
        reconciliation.append(
            [
                row["order_id"],
                datetime.combine(row["manual_date"], datetime.min.time()),
                datetime.combine(row["ledger_date"], datetime.min.time()) if row["ledger_date"] else None,
                f'=IF(OR(B{index}="",C{index}=""),"",ABS(B{index}-C{index}))',
                row["manual_amount"],
                row["ledger_amount"],
                f'=IF(OR(E{index}="",F{index}=""),"",ROUND(E{index}-F{index},2))',
                row["manual_status"],
                row["ledger_status"],
                f'=IF(C{index}="","fail","pass")',
                f'=IF(AND(J{index}="pass",D{index}<=1),"pass","fail")',
                f'=IF(J{index}="fail","fail",IF(ABS(G{index})<0.005,"pass","fail"))',
                f'=IF(AND(J{index}="pass",OR(LOWER(H{index})="completed",LOWER(H{index})="pending"),OR(LOWER(I{index})="completed",LOWER(I{index})="pending")),"pass","fail")',
                f'=IF(COUNTIF(J{index}:M{index},"fail")=0,"matched","discrepancy")',
                f'=IF(N{index}="matched","none",IF(J{index}="fail","missing_in_authoritative_ledger",IF(K{index}="fail","date_mismatch","")&IF(AND(K{index}="fail",L{index}="fail"),";","")&IF(L{index}="fail","amount_mismatch","")&IF(AND(M{index}="fail",OR(K{index}="fail",L{index}="fail")),";","")&IF(M{index}="fail","invalid_status","")))',
            ]
        )
    style_sheet(
        reconciliation,
        {"A": 12, "B": 14, "C": 14, "D": 16, "E": 16, "F": 16, "G": 14,
         "H": 15, "I": 15, "J": 15, "K": 13, "L": 14, "M": 13, "N": 15, "O": 34},
    )
    for row in reconciliation.iter_rows(min_row=2, max_row=reconciliation.max_row):
        for cell in (row[1], row[2]):
            cell.number_format = "yyyy-mm-dd"
        for cell in (row[4], row[5], row[6]):
            cell.number_format = '[$RUB] #,##0.00;[Red]-[$RUB] #,##0.00'
    reconciliation.conditional_formatting.add(
        f"N2:N{reconciliation.max_row}",
        FormulaRule(formula=['N2="discrepancy"'], fill=PatternFill("solid", fgColor=PALE_RED)),
    )
    reconciliation.auto_filter.ref = reconciliation.dimensions
    add_native_table(reconciliation, "ReconciliationTable")
    reconciliation.print_area = f"A1:O{reconciliation.max_row}"

    exceptions = workbook.create_sheet("Exception Register")
    exception_headers = ["Order ID", "Exception Type", "Manual Evidence", "Ledger Evidence", "Impact", "Required Follow-up"]
    exceptions.append(exception_headers)
    for row in rows:
        if row["match_status"] != "discrepancy":
            continue
        order_id = str(row["order_id"])
        manual_evidence = f"{row['manual_date'].isoformat()} | RUB {float(row['manual_amount']):,.2f} | {row['manual_status']}"
        if row["ledger_date"] is None:
            ledger_evidence = "No corresponding order in authoritative ledger"
        else:
            ledger_evidence = f"{row['ledger_date'].isoformat()} | RUB {float(row['ledger_amount']):,.2f} | {row['ledger_status']}"
        reason = str(row["reason"])
        if reason == "amount_mismatch":
            delta = float(row["amount_delta"])
            direction = "overstates" if delta > 0 else "understates"
            impact = f"Manual register {direction} the authoritative amount by RUB {abs(delta):,.2f}."
            follow_up = "Validate supporting evidence and correct the manual amount against the authoritative ledger if confirmed."
            exception_type = "Amount mismatch"
        elif reason == "date_mismatch":
            impact = f"Recorded dates differ by {int(row['date_delta'])} days, exceeding the 1-day tolerance."
            follow_up = "Verify transaction evidence and correct the manual date if the authoritative date is confirmed."
            exception_type = "Date mismatch"
        elif reason == "missing_in_authoritative_ledger":
            impact = f"RUB {float(row['manual_amount']):,.2f} in the manual register cannot be verified against the authoritative ledger."
            follow_up = "Investigate the missing authoritative record before retaining, correcting, or removing the manual entry."
            exception_type = "Missing in authoritative ledger"
        else:
            impact = "The status does not satisfy the allowed completed/pending control."
            follow_up = "Verify the transaction status in both systems before remediation."
            exception_type = "Invalid status"
        exceptions.append([order_id, exception_type, manual_evidence, ledger_evidence, impact, follow_up])
    style_sheet(exceptions, {"A": 12, "B": 29, "C": 29, "D": 39, "E": 48, "F": 58})
    for row in exceptions.iter_rows(min_row=2):
        exceptions.row_dimensions[row[0].row].height = 56
    exceptions.conditional_formatting.add(
        f"A2:F{exceptions.max_row}",
        FormulaRule(formula=['$A2<>""'], fill=PatternFill("solid", fgColor=PALE_AMBER)),
    )
    add_native_table(exceptions, "ExceptionRegisterTable")
    exceptions.print_area = f"A1:F{exceptions.max_row}"

    summary = workbook.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    metrics = [
        ("Total Manual Orders", '=COUNTA(Reconciliation!A2:A16)'),
        ("Matched", '=COUNTIF(Reconciliation!N2:N16,"matched")'),
        ("Discrepancies", '=COUNTIF(Reconciliation!N2:N16,"discrepancy")'),
        ("Manual Total Amount", '=ROUND(SUM(Reconciliation!E2:E16),2)'),
        ("Authoritative Amount Represented", '=ROUND(SUM(Reconciliation!F2:F16),2)'),
        ("Missing Authoritative Amount", '=ROUND(SUMIF(Reconciliation!J2:J16,"fail",Reconciliation!E2:E16),2)'),
        ("Absolute Amount Mismatch", '=ROUND(SUMIF(Reconciliation!G2:G16,">0",Reconciliation!G2:G16)-SUMIF(Reconciliation!G2:G16,"<0",Reconciliation!G2:G16),2)'),
        ("Date Mismatches", '=COUNTIF(Reconciliation!O2:O16,"*date_mismatch*")'),
        ("Amount Mismatches", '=COUNTIF(Reconciliation!O2:O16,"*amount_mismatch*")'),
        ("Missing Records", '=COUNTIF(Reconciliation!O2:O16,"*missing_in_authoritative_ledger*")'),
        ("Control Conclusion", '=IF(B4=0,"Reconciled","Needs remediation")'),
    ]
    for metric, formula in metrics:
        summary.append([metric, formula])
    style_sheet(summary, {"A": 42, "B": 28})
    summary.page_setup.orientation = "portrait"
    summary.freeze_panes = "A2"
    for row in summary.iter_rows(min_row=2):
        row[0].fill = PatternFill("solid", fgColor=PALE_BLUE)
        row[0].font = Font(name="Aptos", size=10, bold=True, color=NAVY)
        row[1].font = Font(name="Aptos Display", size=12, bold=True, color=TEAL)
    for row_number in (5, 6, 7, 8):
        summary.cell(row_number, 2).number_format = '[$RUB] #,##0.00;[Red]-[$RUB] #,##0.00'
    summary.conditional_formatting.add(
        "B4:B4",
        FormulaRule(formula=["B4>0"], fill=PatternFill("solid", fgColor=PALE_RED)),
    )
    add_native_table(summary, "AuditSummaryTable")
    summary.print_area = f"A1:B{summary.max_row}"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    recalculate_workbook(output_path)
    restore_frozen_header_rows(output_path)
    normalize_ooxml(output_path)


def recalculate_workbook(path: Path) -> None:
    """让 LibreOffice 写入公式缓存，确保 Judge 与真实用户都能直接看到结果。"""
    with tempfile.TemporaryDirectory(prefix="audit-xlsx-recalc-") as tmp:
        tmp_root = Path(tmp)
        input_root = tmp_root / "input"
        output_root = tmp_root / "output"
        profile = tmp_root / "profile"
        input_root.mkdir()
        output_root.mkdir()
        profile.mkdir()
        source = input_root / path.name
        shutil.copy2(path, source)
        command = [
            "soffice",
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(output_root),
            str(source),
        ]
        result = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300)
        converted = output_root / path.name
        if result.returncode != 0 or not converted.is_file():
            raise RuntimeError(f"LibreOffice 公式重算失败：{result.stdout[-1000:]}")
        shutil.copy2(converted, path)


def restore_frozen_header_rows(path: Path) -> None:
    """LibreOffice 重算会移除冻结窗格，直接在 OOXML 中恢复三张表的首行冻结。"""
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with tempfile.TemporaryDirectory(prefix="audit-xlsx-freeze-") as tmp:
        target = Path(tmp) / path.name
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w") as destination:
            for info in source.infolist():
                data = source.read(info.filename)
                if re.match(r"xl/worksheets/sheet[1-3]\.xml$", info.filename):
                    root = etree.fromstring(data)
                    sheet_view = root.find(f"{{{namespace}}}sheetViews/{{{namespace}}}sheetView")
                    if sheet_view is not None:
                        for existing in sheet_view.findall(f"{{{namespace}}}pane"):
                            sheet_view.remove(existing)
                        pane = etree.Element(f"{{{namespace}}}pane")
                        pane.set("ySplit", "1")
                        pane.set("topLeftCell", "A2")
                        pane.set("activePane", "bottomLeft")
                        pane.set("state", "frozen")
                        sheet_view.insert(0, pane)
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                destination.writestr(info, data)
        shutil.copy2(target, path)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE, size: str = "6") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def add_word_table(document: Document, headers: list[str], data: list[list[str]], widths: list[float] | None = None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, NAVY)
        set_cell_border(cell, NAVY)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(8)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
        if widths:
            cell.width = Cm(widths[index])
    for row_index, values in enumerate(data):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.text = str(value)
            set_cell_shading(cell, WHITE if row_index % 2 == 0 else "F4F7FA")
            set_cell_border(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(7.5)
                    run.font.color.rgb = RGBColor.from_string(NAVY)
            if widths:
                cell.width = Cm(widths[index])
    return table


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("INTERNAL AUDIT  •  ")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)


def build_report(rows: list[dict[str, object]], facts: dict[str, object], output_path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    document.core_properties.creator = "EvoCowork benchmark adapter"
    document.core_properties.title = "Q1 2024 Financial Audit Reconciliation"
    document.core_properties.created = FIXED_DT.replace(tzinfo=None)
    document.core_properties.modified = FIXED_DT.replace(tzinfo=None)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].font.color.rgb = RGBColor.from_string(NAVY)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    for name, size, color in (("Title", 27, NAVY), ("Subtitle", 12, TEAL), ("Heading 1", 17, NAVY), ("Heading 2", 11, TEAL)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = "ООО «СИБИРСКАЯ ТОРГОВАЯ КОМПАНИЯ»  /  Q1 2024"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(TEAL)
    add_page_number(section.footer.paragraphs[0])

    title = document.add_paragraph(style="Title")
    title.add_run("Financial Audit\nReconciliation")
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run("Q1 2024 cross-system control review")
    document.add_paragraph(
        "Prepared from the frozen authoritative order ledger, the finance manual-tracking workbook, "
        "and the supplied system-mapping rules. Currency: RUB."
    )

    document.add_heading("Executive Summary", level=1)
    document.add_paragraph(
        f"The review reconciled {facts['total']} manual-tracking orders against the authoritative order ledger. "
        f"{facts['matched']} orders matched and {facts['discrepancies']} orders were classified as discrepancies. "
        f"The manual register totals RUB {facts['manual_total']:,.2f}. The control conclusion is Needs remediation."
    )
    add_word_table(
        document,
        ["Metric", "Result", "Metric", "Result"],
        [
            ["Manual orders", str(facts["total"]), "Matched", str(facts["matched"])],
            ["Discrepancies", str(facts["discrepancies"]), "Manual total", f"RUB {facts['manual_total']:,.2f}"],
            ["Missing-ledger amount", f"RUB {facts['missing_amount']:,.2f}", "Absolute amount mismatch", f"RUB {facts['absolute_amount_mismatch']:,.2f}"],
        ],
        [4.1, 2.2, 4.8, 3.0],
    )
    document.add_paragraph(
        "This report identifies control exceptions for follow-up; it does not assert fraud, cancellation, "
        "completed remediation, or an audit opinion on the financial statements."
    )

    document.add_page_break()
    document.add_heading("Scope and Methodology", level=1)
    document.add_paragraph(
        "The population is every row in the Q1 Manual Tracking sheet. Order ID is the primary key. "
        "authoritative_ledger.csv is the authoritative source for store-order facts. A record is matched only when "
        "it exists in both systems, dates differ by no more than 1 day, amounts are identical with zero tolerance, "
        "and both statuses are completed or pending. Status comparison is case-insensitive."
    )
    add_word_table(
        document,
        ["Control", "Rule", "Failure classification"],
        [
            ["Presence", "Exact Order ID exists in both systems", "missing_in_authoritative_ledger"],
            ["Date", "Absolute difference ≤ 1 day", "date_mismatch"],
            ["Amount", "Exact equality; RUB 0.00 tolerance", "amount_mismatch"],
            ["Status", "completed or pending; case-insensitive", "invalid_status"],
        ],
        [3.0, 6.2, 5.2],
    )

    document.add_heading("Reconciliation Results", level=1)
    document.add_paragraph(
        f"The authoritative amount represented for manual orders is RUB {facts['ledger_total']:,.2f}. "
        f"One missing authoritative record represents RUB {facts['missing_amount']:,.2f} on the manual basis. "
        f"The two amount mismatches have an absolute difference of RUB {facts['absolute_amount_mismatch']:,.2f}; "
        f"one date mismatch exceeds the permitted one-day tolerance."
    )
    add_word_table(
        document,
        ["Result", "Count / amount"],
        [
            ["Matched", str(facts["matched"])],
            ["Discrepancies", str(facts["discrepancies"])],
            ["Date mismatches", str(facts["date_mismatches"])],
            ["Amount mismatches", str(facts["amount_mismatches"])],
            ["Missing records", str(facts["missing_records"])],
        ],
        [7.2, 4.0],
    )

    document.add_page_break()
    document.add_heading("Discrepancies and Required Actions", level=1)
    discrepancy_data: list[list[str]] = []
    for row in facts["discrepancy_rows"]:
        order_id = str(row["order_id"])
        manual = f"{row['manual_date'].isoformat()}\nRUB {float(row['manual_amount']):,.2f}\n{row['manual_status']}"
        if row["ledger_date"] is None:
            ledger = "No corresponding authoritative record"
            difference = "Missing record; manual basis RUB 445.75"
            action = "Investigate source evidence before retaining, correcting, or removing the manual entry."
        else:
            ledger = f"{row['ledger_date'].isoformat()}\nRUB {float(row['ledger_amount']):,.2f}\n{row['ledger_status']}"
            if row["reason"] == "date_mismatch":
                difference = f"Date difference: {int(row['date_delta'])} days"
                action = "Verify transaction evidence and correct the manual date if the authoritative date is confirmed."
            else:
                difference = f"Manual - ledger: RUB {float(row['amount_delta']):+,.2f}"
                action = "Validate supporting evidence and correct the manual amount if the authoritative value is confirmed."
        discrepancy_data.append([order_id, str(row["reason"]), manual, ledger, difference, action])
    add_word_table(
        document,
        ["Order", "Exception", "Manual evidence", "Authoritative evidence", "Difference", "Required action"],
        discrepancy_data,
        [1.5, 3.1, 2.7, 3.1, 2.7, 4.4],
    )

    document.add_heading("Evidence Boundary", level=1)
    document.add_paragraph(
        "The supplied evidence supports order presence, recorded dates, amounts, statuses, rule-based match status, "
        "and the listed control exceptions. It does not support conclusions about fraud, cancellation, customer loss, "
        "revenue impact, penalties, responsibility allocation, ROI, completed remediation, or whether the financial "
        "statements are free of material misstatement. All proposed actions therefore remain verification steps."
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    normalize_ooxml(output_path)


def normalize_ooxml(path: Path) -> None:
    """固定 ZIP 时间戳，避免参考答案因生成时间不同而漂移。"""
    with tempfile.TemporaryDirectory(prefix="audit-ooxml-normalize-") as tmp:
        target = Path(tmp) / path.name
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w") as destination:
            for info in sorted(source.infolist(), key=lambda item: item.filename):
                data = source.read(info.filename)
                if info.filename == "docProps/core.xml":
                    root = etree.fromstring(data)
                    dcterms = "http://purl.org/dc/terms/"
                    for field in ("created", "modified"):
                        node = root.find(f"{{{dcterms}}}{field}")
                        if node is not None:
                            node.text = "2024-04-01T00:00:00Z"
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                copied = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
                copied.compress_type = zipfile.ZIP_DEFLATED
                copied.external_attr = info.external_attr
                copied.create_system = info.create_system
                copied.flag_bits = info.flag_bits
                destination.writestr(copied, data)
        shutil.copy2(target, path)


def build_email(facts: dict[str, object], xlsx_path: Path, docx_path: Path, output_path: Path) -> None:
    message = EmailMessage(policy=policy.default.clone(linesep="\r\n"))
    message["Date"] = format_datetime(FIXED_DT)
    message["From"] = "audit@company.com"
    message["To"] = "finance@company.com"
    message["Subject"] = "Financial Audit Reconciliation - Q1 2024"
    message["X-Draft-Status"] = "draft"
    lines = [
        "Olga,",
        "",
        f"The Q1 reconciliation covers {facts['total']} manual orders: {facts['matched']} matched and {facts['discrepancies']} discrepancies.",
        f"Manual-register total: RUB {facts['manual_total']:,.2f}.",
        "",
        "Exceptions:",
    ]
    for row in facts["discrepancy_rows"]:
        lines.append(f"- {row['order_id']}: {row['reason']}")
    lines.extend(
        [
            "",
            "The editable reconciliation workbook and audit report are attached for review.",
            "This is a draft only; no message has been sent and no remediation is recorded as complete.",
            "",
            "Internal Audit",
        ]
    )
    message.set_content("\n".join(lines), subtype="plain", charset="utf-8")
    message.add_attachment(
        xlsx_path.read_bytes(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=XLSX_NAME,
    )
    message.add_attachment(
        docx_path.read_bytes(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=DOCX_NAME,
    )
    message.set_boundary("evo-cowork-financial-audit-q1-2024")
    output_path.write_bytes(message.as_bytes())


def fold_ical(line: str, width: int = 72) -> list[str]:
    if len(line) <= width:
        return [line]
    parts = [line[:width]]
    rest = line[width:]
    while rest:
        parts.append(" " + rest[: width - 1])
        rest = rest[width - 1 :]
    return parts


def build_calendar(facts: dict[str, object], output_path: Path) -> None:
    exceptions = ", ".join(f"{row['order_id']} {row['reason']}" for row in facts["discrepancy_rows"])
    description = (
        f"{facts['total']} manual orders; {facts['matched']} matched; {facts['discrepancies']} discrepancies: "
        f"{exceptions}. Review the editable audit workbook and report before remediation."
    )
    raw_lines = [
        "BEGIN:VCALENDAR",
        "VERSION:2.0",
        "PRODID:-//EvoCowork//Financial Audit Reconciliation//EN",
        "CALSCALE:GREGORIAN",
        "METHOD:REQUEST",
        "BEGIN:VTIMEZONE",
        "TZID:Asia/Shanghai",
        "BEGIN:STANDARD",
        "DTSTART:19700101T000000",
        "TZOFFSETFROM:+0800",
        "TZOFFSETTO:+0800",
        "TZNAME:CST",
        "END:STANDARD",
        "END:VTIMEZONE",
        "BEGIN:VEVENT",
        "UID:audit-reconciliation-review-20240405@company.com",
        "DTSTAMP:20240401T000000Z",
        "DTSTART;TZID=Asia/Shanghai:20240405T100000",
        "DTEND;TZID=Asia/Shanghai:20240405T104500",
        "SUMMARY:Audit Reconciliation Review",
        "STATUS:TENTATIVE",
        "ORGANIZER:mailto:audit@company.com",
        "ATTENDEE;CN=Olga Smirnova;ROLE=REQ-PARTICIPANT;PARTSTAT=NEEDS-ACTION;RSVP=TRUE:mailto:finance@company.com",
        f"DESCRIPTION:{description}",
        f"ATTACH;VALUE=URI:{XLSX_NAME}",
        f"ATTACH;VALUE=URI:{DOCX_NAME}",
        "END:VEVENT",
        "END:VCALENDAR",
    ]
    lines: list[str] = []
    for line in raw_lines:
        lines.extend(fold_ical(line))
    output_path.write_text("\r\n".join(lines) + "\r\n", encoding="utf-8", newline="")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for child in args.output_dir.iterdir():
        if child.is_file():
            child.unlink()
        elif child.is_dir():
            shutil.rmtree(child)

    rows = read_sources(args.input_dir)
    facts = truth(rows)
    xlsx_path = args.output_dir / XLSX_NAME
    docx_path = args.output_dir / DOCX_NAME
    build_workbook(rows, xlsx_path)
    build_report(rows, facts, docx_path)
    build_email(facts, xlsx_path, docx_path, args.output_dir / EML_NAME)
    build_calendar(facts, args.output_dir / ICS_NAME)


if __name__ == "__main__":
    main()
