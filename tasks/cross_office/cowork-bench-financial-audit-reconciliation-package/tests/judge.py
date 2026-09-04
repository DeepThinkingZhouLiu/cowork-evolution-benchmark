#!/usr/bin/env python3
"""Q1 跨系统财务审计包的 v2 连续评分入口；v1 原样保存在 judge_v1.py。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import re
import sys
from email import policy
from email.parser import BytesParser
from pathlib import Path
from typing import Iterable

from docx import Document
from openpyxl import load_workbook


TASK_ID = "cowork-evo/cowork-bench-financial-audit-reconciliation-package"
RUBRIC_VERSION = "cowork-bench-financial-audit-reconciliation-v0.2"
JUDGE_VERSION = "cowork-bench-financial-audit-reconciliation-criterion-partial-v2.0"
XLSX = "Audit_Reconciliation_Report.xlsx"
DOCX = "Audit_Report.docx"
EML = "Financial_Audit_Reconciliation.eml"
ICS = "Audit_Reconciliation_Review.ics"
REQUIRED = (XLSX, DOCX, EML, ICS)
POSITIVE_WEIGHTS = {
    "R001": 5, "R002": 15, "R003": 15, "R004": 8, "R005": 10, "R006": 8,
    "R007": 8, "R008": 7, "R009": 7, "R010": 7, "R011": 5, "R012": 5,
}
PENALTY_WEIGHTS = {"R013": -30, "R014": -30, "R015": -25, "R016": -25, "R017": -25}
SUMMARY_FORMULA_TOKENS = (
    ("COUNTA", "RECONCILIATION", "A2:A16"),
    ("COUNTIF", "RECONCILIATION", "N2:N16", "MATCHED"),
    ("COUNTIF", "RECONCILIATION", "N2:N16", "DISCREPANCY"),
    ("ROUND", "SUM", "RECONCILIATION", "E2:E16"),
    ("ROUND", "SUM", "RECONCILIATION", "F2:F16"),
    ("ROUND", "SUMIF", "RECONCILIATION", "J2:J16", "FAIL", "E2:E16"),
    ("ROUND", "SUMIF", "RECONCILIATION", "G2:G16", ">0", "<0"),
    ("COUNTIF", "RECONCILIATION", "O2:O16", "DATE_MISMATCH"),
    ("COUNTIF", "RECONCILIATION", "O2:O16", "AMOUNT_MISMATCH"),
    ("COUNTIF", "RECONCILIATION", "O2:O16", "MISSING_IN_AUTHORITATIVE_LEDGER"),
    ("IF", "B4", "RECONCILED", "NEEDS REMEDIATION"),
)


def _load_v1():
    path = Path(__file__).with_name("judge_v1.py")
    if str(path.parent) not in sys.path:
        sys.path.insert(0, str(path.parent))
    spec = importlib.util.spec_from_file_location("financial_audit_judge_v1", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v1 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


V1 = _load_v1()


def _sha(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _criterion(key: str, score: float, evidence: object, earned: float = 0, total: float = 1) -> dict[str, object]:
    return {
        "criterion_id": key,
        "score": round(max(0.0, min(1.0, float(score))), 6),
        "evidence": evidence,
        "method": "deterministic",
        "raw": {"earned_units": round(earned, 6), "total_units": round(total, 6)},
    }


def _ratio(findings: Iterable[tuple[bool, str]]) -> tuple[float, list[str], int, int]:
    rows = list(findings)
    earned = sum(bool(ok) for ok, _ in rows)
    return earned / max(1, len(rows)), [label for ok, label in rows if not ok], earned, len(rows)


def _norm(value: object) -> str:
    return V1.normalize(value)


def _lower(value: object) -> str:
    return V1.lower(value)


def _formula(cell, *tokens: str) -> bool:
    text = _norm(cell.value).upper()
    return cell.data_type == "f" and all(token.upper() in text for token in tokens)


def _same(actual: object, expected: object) -> bool:
    if isinstance(expected, float):
        return V1.close_number(actual, expected)
    if expected is None:
        return actual in (None, "")
    return _norm(actual) == _norm(expected)


def _message(path: Path):
    return BytesParser(policy=policy.default).parsebytes(path.read_bytes())


def _mail_body(message) -> str:
    body = ""
    try:
        if message.is_multipart():
            for part in message.walk():
                if part.get_content_type() == "text/plain" and part.get_content_disposition() != "attachment":
                    body += part.get_content()
        else:
            body = message.get_content()
    except Exception:
        body = ""
    return body


def _core(artifact: Path) -> tuple[bool, list[str], list[str]]:
    errors: list[str] = []
    hashes: list[str] = []
    source_ok, source_errors = V1.source_integrity()
    if not source_ok:
        errors.append(f"冻结输入哈希失败：{source_errors}")
    if not artifact.is_dir():
        return False, ["输出目录不存在", *errors], hashes

    source_hashes = set(V1.SOURCE_HASHES.values())
    for name in REQUIRED:
        path = artifact / name
        if not path.is_file() or path.stat().st_size == 0:
            errors.append(f"核心交付缺失或为空：{name}")
            continue
        digest = _sha(path)
        hashes.append(digest)
        if digest in source_hashes:
            errors.append(f"核心交付只是原样输入：{name}")
            continue
        try:
            if name == XLSX:
                workbook = load_workbook(path, read_only=True, data_only=False)
                if not workbook.sheetnames or all(sheet.max_row == 0 for sheet in workbook.worksheets):
                    raise ValueError("工作簿没有内容")
                workbook.close()
            elif name == DOCX:
                document = Document(path)
                if not document.paragraphs and not document.tables:
                    raise ValueError("Word 没有内容")
            elif name == EML:
                message = _message(path)
                if not list(message.items()) or not message.get_payload():
                    raise ValueError("邮件没有可解析头部或正文")
            else:
                text = path.read_text(encoding="utf-8")
                if "BEGIN:VCALENDAR" not in text or "BEGIN:VEVENT" not in text:
                    raise ValueError("不是可解析的 VCALENDAR/VEVENT")
        except Exception as exc:  # noqa: BLE001
            errors.append(f"核心交付损坏或不可解析：{name}: {type(exc).__name__}: {exc}")
    return not errors, errors, hashes


def _fatal(errors: list[str], hashes: list[str]) -> dict[str, object]:
    criteria = [_criterion(key, 0, "；".join(errors), 0, 0) for key in (*POSITIVE_WEIGHTS, *PENALTY_WEIGHTS)]
    return {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes or ["0" * 64], "criterion_results": criteria,
        "reward": 0.0, "passed": False, "hurdle_failures": ["R001"], "penalties": [],
    }


def _reconciliation_formula_ok(cell, row: int, column: int) -> bool:
    rules = {
        4: ("ABS", f"B{row}", f"C{row}"),
        7: ("ROUND", f"E{row}", f"F{row}"),
        10: (f"C{row}", "PASS", "FAIL"),
        11: (f"J{row}", f"D{row}", "<=1", "PASS", "FAIL"),
        12: (f"J{row}", f"G{row}", "ABS", "PASS", "FAIL"),
        13: (f"J{row}", f"H{row}", f"I{row}", "LOWER", "COMPLETED", "PENDING"),
        14: (f"J{row}:M{row}", "MATCHED", "DISCREPANCY"),
        15: (f"N{row}", f"J{row}", f"K{row}", f"L{row}", f"M{row}",
             "MISSING_IN_AUTHORITATIVE_LEDGER", "DATE_MISMATCH", "AMOUNT_MISMATCH", "INVALID_STATUS"),
    }
    return _formula(cell, *rules[column])


def _workbook(path: Path, rows: list[dict[str, object]], facts: dict[str, object]):
    source_findings: list[tuple[bool, str]] = []
    formula_findings: list[tuple[bool, str]] = []
    summary_findings: list[tuple[bool, str]] = []
    exception_findings: list[tuple[bool, str]] = []
    native_findings: list[tuple[bool, str]] = []
    formula_total = 131
    formula_count = 0
    delivered_ids: list[str] = []
    exception_ids: set[str] = set()
    workbook = load_workbook(path, data_only=False)

    expected_sheets = ["Reconciliation", "Exception Register", "Summary"]
    source_findings.append((workbook.sheetnames == expected_sheets, f"工作表顺序/集合错误：{workbook.sheetnames}"))
    if "Reconciliation" not in workbook.sheetnames:
        formula_findings.append((False, "缺少 Reconciliation"))
        summary_findings.append((False, "缺少 Summary"))
        exception_findings.append((False, "缺少 Exception Register"))
        return (_ratio(source_findings), _ratio(formula_findings), _ratio(summary_findings),
                _ratio(exception_findings), _ratio(native_findings), delivered_ids, exception_ids, formula_count, formula_total, workbook)

    recon = workbook["Reconciliation"]
    headers = [_norm(recon.cell(1, column).value) for column in range(1, 16)]
    for index, expected in enumerate(V1.RECON_HEADERS, start=1):
        source_findings.append((index <= len(headers) and headers[index - 1] == expected, f"Reconciliation 表头 {index} 错误"))
    source_findings.append((recon.max_row == 16 and recon.max_column == 15, f"Reconciliation 形状错误：{recon.max_row}x{recon.max_column}"))
    for row_number, expected in enumerate(rows, start=2):
        delivered_ids.append(_norm(recon.cell(row_number, 1).value))
        source_findings.append((_norm(recon.cell(row_number, 1).value) == expected["order_id"], f"A{row_number} 订单错误"))
        expected_static = {
            2: expected["manual_date"], 3: expected["ledger_date"], 5: expected["manual_amount"],
            6: expected["ledger_amount"], 8: expected["manual_status"], 9: expected["ledger_status"],
        }
        for column, wanted in expected_static.items():
            actual = recon.cell(row_number, column).value
            if column in (2, 3):
                ok = V1.iso_date(actual) == wanted
            elif column in (5, 6):
                ok = actual in (None, "") if wanted is None else V1.close_number(actual, float(wanted))
            else:
                ok = actual in (None, "") if wanted is None else _lower(actual) == _lower(wanted)
            source_findings.append((ok, f"{recon.cell(row_number, column).coordinate} 来源值错误"))
        for column in (4, 7, 10, 11, 12, 13, 14, 15):
            cell = recon.cell(row_number, column)
            ok = _reconciliation_formula_ok(cell, row_number, column)
            formula_findings.append((ok, f"{cell.coordinate} 公式语义/行引用错误"))
            formula_count += int(cell.data_type == "f")

    if "Summary" in workbook.sheetnames:
        summary = workbook["Summary"]
        summary_findings.extend([
            (_norm(summary["A1"].value) == "Metric", "Summary!A1 错误"),
            (_norm(summary["B1"].value) == "Value", "Summary!B1 错误"),
        ])
        for index, metric in enumerate(V1.SUMMARY_METRICS, start=2):
            summary_findings.append((_norm(summary.cell(index, 1).value) == metric, f"Summary!A{index} 指标错误"))
            cell = summary.cell(index, 2)
            ok = _formula(cell, *SUMMARY_FORMULA_TOKENS[index - 2])
            summary_findings.append((ok, f"Summary!B{index} 公式语义错误"))
            formula_count += int(cell.data_type == "f")
    else:
        summary_findings.append((False, "缺少 Summary"))

    if "Exception Register" in workbook.sheetnames:
        sheet = workbook["Exception Register"]
        for index, expected_header in enumerate(V1.EXCEPTION_HEADERS, start=1):
            exception_findings.append((_norm(sheet.cell(1, index).value) == expected_header, f"Exception Register 表头 {index} 错误"))
        records = {
            _norm(sheet.cell(row, 1).value): [_norm(sheet.cell(row, column).value) for column in range(1, 7)]
            for row in range(2, sheet.max_row + 1) if _norm(sheet.cell(row, 1).value)
        }
        exception_ids = set(records)
        expected_ids = set(facts["discrepancy_ids"])
        exception_findings.append((exception_ids == expected_ids, f"异常订单集合错误：{sorted(exception_ids)}"))
        for expected in facts["discrepancy_rows"]:
            order_id = str(expected["order_id"])
            delivered = records.get(order_id, [""] * 6)
            blob = _lower(" | ".join(delivered))
            reason = str(expected["reason"])
            exception_findings.append((V1.reason_present(blob, reason), f"{order_id} 异常类型错误"))
            exception_findings.append((expected["manual_date"] in blob and V1.has_money(blob, float(expected["manual_amount"])), f"{order_id} 手工证据不全"))
            if expected["ledger_date"] is None:
                ledger_ok = ("missing" in blob or "no corresponding" in blob or "缺失" in blob) and V1.has_money(blob, float(expected["manual_amount"]))
            else:
                ledger_ok = expected["ledger_date"] in blob and V1.has_money(blob, float(expected["ledger_amount"]))
            exception_findings.append((ledger_ok, f"{order_id} 权威证据不全"))
            if reason == "amount_mismatch":
                impact_ok = V1.has_money(delivered[4], abs(float(expected["amount_delta"])))
            elif reason == "date_mismatch":
                impact_ok = str(expected["date_delta"]) in delivered[4] and ("day" in _lower(delivered[4]) or "天" in delivered[4])
            else:
                impact_ok = V1.has_money(delivered[4], float(expected["manual_amount"]))
            exception_findings.append((impact_ok, f"{order_id} 差异影响不准确"))
            action = _lower(delivered[5])
            action_ok = len(action) >= 20 and any(token in action for token in ("verify", "validate", "investigate", "correct", "核验", "调查", "更正"))
            exception_findings.append((action_ok, f"{order_id} 后续动作不充分"))
    else:
        exception_findings.append((False, "缺少 Exception Register"))

    native_findings.extend([
        (sum(len(sheet.tables) for sheet in workbook.worksheets) >= 3, "Excel 原生表少于 3"),
        (all(name in workbook.sheetnames and workbook[name].freeze_panes == "A2" for name in expected_sheets), "Excel 未全部冻结表头"),
        (all(name in workbook.sheetnames and bool(workbook[name].auto_filter.ref) for name in expected_sheets), "Excel 未全部启用筛选"),
        (sum(len(sheet.conditional_formatting) for sheet in workbook.worksheets) >= 2, "Excel 条件格式不足"),
        (all(recon.cell(row, column).number_format not in ("General", "") for row in range(2, 17) for column in (2, 3, 5, 6, 7)), "日期/金额格式不完整"),
    ])
    return (_ratio(source_findings), _ratio(formula_findings), _ratio(summary_findings),
            _ratio(exception_findings), _ratio(native_findings), delivered_ids, exception_ids,
            formula_count, formula_total, workbook)


def _word(path: Path, facts: dict[str, object]):
    document = Document(path)
    text = V1.document_text(document)
    blob = _lower(text)
    headings = [_norm(paragraph.text) for paragraph in document.paragraphs if str(paragraph.style.name).startswith("Heading 1")]
    summary = [(heading in headings, f"Word 缺一级标题 {heading}") for heading in V1.WORD_HEADINGS]
    summary.extend([
        (V1.count_with_label(text, "orders", int(facts["total"])), "Word 总订单数错误"),
        (V1.count_with_label(text, "matched", int(facts["matched"])), "Word matched 数错误"),
        (V1.count_with_label(text, "discrepancies", int(facts["discrepancies"])), "Word discrepancy 数错误"),
        (V1.has_money(text, float(facts["manual_total"])), "Word 手工总金额错误"),
        ("authoritative" in blob and ("source" in blob or "ledger" in blob), "Word 未声明权威来源"),
        ("1 day" in blob or "one-day" in blob or "1-day" in blob, "Word 未声明 1 天日期容差"),
        ("zero tolerance" in blob or "0.00 tolerance" in blob or "零容差" in blob, "Word 未声明金额零容差"),
        ("case-insensitive" in blob or "ignore case" in blob or "忽略大小写" in blob, "Word 未声明状态大小写规则"),
        ("completed" in blob and "pending" in blob, "Word 未声明允许状态"),
    ])
    records: dict[str, str] = {}
    for table in document.tables:
        for row in table.rows[1:]:
            row_text = " | ".join(cell.text for cell in row.cells)
            for order_id in facts["discrepancy_ids"]:
                if order_id in row_text:
                    records[order_id] = row_text
    exceptions: list[tuple[bool, str]] = []
    for expected in facts["discrepancy_rows"]:
        order_id = str(expected["order_id"])
        row_text = records.get(order_id, "")
        exceptions.append((bool(row_text), f"Word 缺异常行 {order_id}"))
        exceptions.append((V1.reason_present(row_text, str(expected["reason"])), f"Word {order_id} 异常类型错误"))
        exceptions.append((expected["manual_date"] in row_text and V1.has_money(row_text, float(expected["manual_amount"])), f"Word {order_id} 手工证据错误"))
        if expected["ledger_date"] is None:
            ledger_ok = ("missing" in _lower(row_text) or "no corresponding" in _lower(row_text) or "缺失" in row_text) and V1.has_money(row_text, float(expected["manual_amount"]))
        else:
            ledger_ok = expected["ledger_date"] in row_text and V1.has_money(row_text, float(expected["ledger_amount"]))
        exceptions.append((ledger_ok, f"Word {order_id} 权威证据错误"))
        if expected["reason"] == "amount_mismatch":
            difference_ok = V1.has_money(row_text, abs(float(expected["amount_delta"])))
        elif expected["reason"] == "date_mismatch":
            difference_ok = str(expected["date_delta"]) in row_text
        else:
            difference_ok = V1.has_money(row_text, float(expected["manual_amount"]))
        exceptions.append((difference_ok, f"Word {order_id} 差异错误"))
        action_ok = any(token in _lower(row_text) for token in ("verify", "validate", "investigate", "correct", "核验", "调查", "更正"))
        exceptions.append((action_ok, f"Word {order_id} 缺谨慎后续动作"))
    native = [
        (len(document.tables) >= 4, "Word 原生表少于 4"),
        (len(headings) >= 5, "Word 一级标题样式不足"),
    ]
    return _ratio(summary), _ratio(exceptions), text, document, native


def _email(path: Path, artifact: Path, facts: dict[str, object]):
    message = _message(path)
    body = _mail_body(message)
    recipients = []
    if message.get("To"):
        try:
            recipients = [address.addr_spec.lower() for address in message["To"].addresses]
        except Exception:
            recipients = [_lower(message.get("To"))]
    content = [
        (recipients == ["finance@company.com"], f"邮件收件人错误：{recipients}"),
        (_norm(message.get("Subject")) == "Financial Audit Reconciliation - Q1 2024", "邮件主题错误"),
        (_lower(message.get("X-Draft-Status")) == "draft", "邮件未标记 draft"),
        (V1.count_with_label(body, "orders", int(facts["total"])), "邮件总订单数错误"),
        (V1.count_with_label(body, "matched", int(facts["matched"])), "邮件 matched 数错误"),
        (V1.count_with_label(body, "discrepancies", int(facts["discrepancies"])), "邮件 discrepancy 数错误"),
        (V1.has_money(body, float(facts["manual_total"])), "邮件手工总金额错误"),
    ]
    for expected in facts["discrepancy_rows"]:
        content.append((str(expected["order_id"]) in body, f"邮件缺 {expected['order_id']}"))
        content.append((V1.reason_present(body, str(expected["reason"])), f"邮件缺 {expected['order_id']} 类型"))

    attachments: dict[str, tuple[bytes, str]] = {}
    for part in message.iter_attachments():
        name = part.get_filename()
        if name:
            attachments[name] = (part.get_payload(decode=True) or b"", part.get_content_type())
    attachment_findings = [(set(attachments) == {XLSX, DOCX}, f"邮件附件集合错误：{sorted(attachments)}")]
    expected_types = {
        XLSX: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        DOCX: "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }
    for name in (XLSX, DOCX):
        payload, mime = attachments.get(name, (b"", ""))
        attachment_findings.extend([
            (bool(payload), f"邮件附件 {name} 缺失/为空"),
            (mime == expected_types[name], f"邮件附件 {name} MIME 错误"),
            ((artifact / name).is_file() and payload == (artifact / name).read_bytes(), f"邮件附件 {name} 与目录文件不一致"),
        ])
    return _ratio(content), _ratio(attachment_findings), body, attachments


def _unfold_ics(text: str) -> tuple[list[str], dict[str, list[str]]]:
    lines = V1.unfold_ics(text)
    properties: dict[str, list[str]] = {}
    for line in lines:
        if ":" not in line:
            continue
        key, value = line.split(":", 1)
        properties.setdefault(key, []).append(value)
    return lines, properties


def _calendar(path: Path, facts: dict[str, object]):
    text = path.read_text(encoding="utf-8")
    lines, properties = _unfold_ics(text)
    description = " ".join(value for key, values in properties.items() if key == "DESCRIPTION" for value in values)
    attachments = [value for key, values in properties.items() if key.startswith("ATTACH") for value in values]
    attendees = [value.lower() for key, values in properties.items() if key.startswith("ATTENDEE") for value in values]
    findings = [
        ("BEGIN:VCALENDAR" in lines and "END:VCALENDAR" in lines, "ICS 日历外壳错误"),
        ("BEGIN:VEVENT" in lines and "END:VEVENT" in lines, "ICS 事件外壳错误"),
        (properties.get("TZID") == ["Asia/Shanghai"], "ICS VTIMEZONE 错误"),
        (properties.get("DTSTART;TZID=Asia/Shanghai") == ["20240405T100000"], "ICS 开始时间错误"),
        (properties.get("DTEND;TZID=Asia/Shanghai") == ["20240405T104500"], "ICS 结束时间错误"),
        (properties.get("SUMMARY") == ["Audit Reconciliation Review"], "ICS 标题错误"),
        (properties.get("STATUS") == ["TENTATIVE"], "ICS 状态错误"),
        (attendees == ["mailto:finance@company.com"], "ICS 参会人错误"),
        (V1.count_with_label(description, "orders", int(facts["total"])), "ICS 总订单数错误"),
        (V1.count_with_label(description, "matched", int(facts["matched"])), "ICS matched 数错误"),
        (V1.count_with_label(description, "discrepancies", int(facts["discrepancies"])), "ICS discrepancy 数错误"),
    ]
    for expected in facts["discrepancy_rows"]:
        findings.append((str(expected["order_id"]) in description, f"ICS 缺 {expected['order_id']}"))
        findings.append((V1.reason_present(description, str(expected["reason"])), f"ICS 缺 {expected['order_id']} 类型"))
    findings.extend([
        (set(attachments) == {XLSX, DOCX}, f"ICS 附件 URI 错误：{attachments}"),
        (not any(marker in _lower(text) for marker in ("file://", "/app/", "/home/", "c:\\")), "ICS 出现绝对路径"),
    ])
    return _ratio(findings), text, description, attachments


def _render_findings(artifact: Path) -> list[tuple[bool, str]]:
    findings = []
    for name, minimum, maximum in ((XLSX, 3, 6), (DOCX, 2, 6)):
        pages, text, log = V1.render_office(artifact / name)
        findings.append((minimum <= pages <= maximum, f"{name} LibreOffice 渲染页数异常：{pages}; {log[-300:]}"))
        required = ["WC-013"] if name == XLSX else [*V1.WORD_HEADINGS, *V1.compute_truth()[1]["discrepancy_ids"]]
        findings.append((all(token in text for token in required), f"{name} 渲染后缺核心可见文本"))
    return findings


def _professional(artifact: Path, workbook, document, native):
    package_findings = []
    for name in (XLSX, DOCX):
        ok, evidence = V1.ooxml_package(artifact / name)
        package_findings.extend([
            (ok, f"{name} OOXML 关系/宏异常：{evidence}"),
            (not evidence.get("external"), f"{name} 含外部关系"),
        ])
    media = []
    for name in (XLSX, DOCX):
        _, evidence = V1.ooxml_package(artifact / name)
        media.extend(evidence.get("media", []))
    findings = [*native, (not media, f"Office 内容被图片化：{media}"), *_render_findings(artifact)]
    return _ratio([*package_findings, *findings])


def _identity_penalty(delivered_ids: list[str], expected_ids: list[str]):
    expected = set(expected_ids)
    delivered = [value for value in delivered_ids if value]
    missing = expected - set(delivered)
    extra = set(delivered) - expected
    duplicates = len(delivered) - len(set(delivered))
    affected = len(missing) + len(extra) + max(0, duplicates)
    severity = min(1.0, affected / max(1, len(expected_ids)))
    labels = [] if not affected else [f"订单身份缺失={sorted(missing)}，额外={sorted(extra)}，重复={duplicates}"]
    return severity, labels, len(expected_ids)


def _summary_conflict(text: str, facts: dict[str, object]) -> bool:
    normalized = _lower(text)
    patterns = (
        ((r"\b(\d+)\s+manual(?:[- ]tracking)?\s+orders\b",), int(facts["total"])),
        ((r"\b(\d+)\s+orders\s+matched\b", r"manual orders\s*[:;]\s*(\d+)\s+matched\b", r"(?:^|;)\s*(\d+)\s+matched\b"), int(facts["matched"])),
        ((r"\b(\d+)\s+orders\s+(?:were\s+classified\s+as\s+)?discrepanc(?:y|ies)\b", r"matched\s+and\s+(\d+)\s+discrepanc(?:y|ies)\b", r"(?:^|;)\s*(\d+)\s+discrepanc(?:y|ies)\b"), int(facts["discrepancies"])),
    )
    for alternatives, expected in patterns:
        for pattern in alternatives:
            for match in re.finditer(pattern, normalized):
                if int(match.group(1)) != expected:
                    return True
    return False


def _cross_penalty(artifact: Path, attachments, word_text: str, mail_body: str, calendar_description: str, facts):
    issues: list[tuple[float, str]] = []
    expected_names = {XLSX, DOCX}
    if set(attachments) != expected_names:
        issues.append((1.0, "邮件附件缺失、额外或命名错误"))
    else:
        stale = [name for name in expected_names if attachments[name][0] != (artifact / name).read_bytes()]
        if stale:
            issues.append((1.0, f"邮件附件与目录文件不一致：{stale}"))
    carriers = [("Word", word_text), ("EML", mail_body), ("ICS", calendar_description)]
    conflicting = [name for name, text in carriers if _summary_conflict(text, facts)]
    if conflicting:
        issues.append((min(0.75, 0.25 * len(conflicting)), f"载体汇总发生冲突：{conflicting}"))
    severity = max((value for value, _ in issues), default=0.0)
    return severity, [label for _, label in issues], 1.0


def _formula_penalty(formula_count: int, formula_total: int):
    coverage = formula_count / max(1, formula_total)
    severity = 0.0 if coverage >= 0.5 else min(1.0, (0.5 - coverage) / 0.5)
    return severity, ([] if severity == 0 else [f"原生公式覆盖率仅 {coverage:.1%}"]), 1.0


def _boundary_penalty(text: str):
    findings = V1.unsupported_claims(text)
    severity = min(1.0, len(findings) / 2)
    return severity, findings, 1.0


def _safe_penalty(artifact: Path, combined_text: str):
    delivered = {path.name for path in artifact.iterdir()}
    extras = sorted(delivered - set(REQUIRED) - {"__pycache__", "solve.py", "solve.sh"})
    suspicious = [marker for marker in V1.SUSPICIOUS_MARKERS if marker in _lower(combined_text)]
    unsafe = []
    for name in (XLSX, DOCX):
        _, evidence = V1.ooxml_package(artifact / name)
        if evidence.get("macros") or evidence.get("external"):
            unsafe.append(name)
    issues: list[tuple[float, str]] = []
    if unsafe:
        issues.append((1.0, f"宏或危险外链：{unsafe}"))
    if suspicious:
        issues.append((0.8, f"绝对路径/Golden/Judge 泄漏：{suspicious}"))
    if extras:
        issues.append((0.4, f"额外输出：{extras}"))
    severity = max((value for value, _ in issues), default=0.0)
    return severity, [label for _, label in issues], 1.0


def evaluate(artifact: Path) -> dict[str, object]:
    core_ok, core_errors, hashes = _core(artifact)
    if not core_ok:
        return _fatal(core_errors, hashes)

    rows, facts = V1.compute_truth()
    (source, formulas, summary, exceptions, native, delivered_ids, _exception_ids,
     formula_count, formula_total, workbook) = _workbook(artifact / XLSX, rows, facts)
    word_summary, word_exceptions, word_text, document, word_native = _word(artifact / DOCX, facts)
    email_content, email_attachments, mail_body, attachments = _email(artifact / EML, artifact, facts)
    calendar, calendar_text, calendar_description, _ics_attachments = _calendar(artifact / ICS, facts)

    cross_findings = [
        (V1.count_with_label(text, "orders", int(facts["total"])), f"{name} 总订单数未保持一致")
        for name, text in (("Word", word_text), ("EML", mail_body), ("ICS", calendar_description))
    ]
    cross_findings.extend([
        (V1.count_with_label(text, "matched", int(facts["matched"])), f"{name} matched 数未保持一致")
        for name, text in (("Word", word_text), ("EML", mail_body), ("ICS", calendar_description))
    ])
    cross_findings.extend([
        (V1.count_with_label(text, "discrepancies", int(facts["discrepancies"])), f"{name} discrepancy 数未保持一致")
        for name, text in (("Word", word_text), ("EML", mail_body), ("ICS", calendar_description))
    ])
    for expected in facts["discrepancy_rows"]:
        for name, text in (("Word", word_text), ("EML", mail_body), ("ICS", calendar_description)):
            cross_findings.append((str(expected["order_id"]) in text and V1.reason_present(text, str(expected["reason"])),
                                   f"{name} 未一致传播 {expected['order_id']} {expected['reason']}"))
    cross = _ratio(cross_findings)
    professional = _professional(artifact, workbook, document, word_native)
    # native 已是一个 ratio；把其失败项和通过数还原为逐项结果，避免丢失 Excel 原子证据。
    native_score, native_failures, native_earned, native_total = native
    professional_score, professional_failures, professional_earned, professional_total = professional
    combined_professional = (
        (native_earned + professional_earned) / max(1, native_total + professional_total),
        [*native_failures, *professional_failures],
        native_earned + professional_earned,
        native_total + professional_total,
    )

    groups = {
        "R002": source, "R003": formulas, "R004": summary, "R005": exceptions,
        "R006": word_summary, "R007": word_exceptions, "R008": email_content,
        "R009": email_attachments, "R010": calendar, "R011": cross, "R012": combined_professional,
    }
    results = [_criterion("R001", 1.0, "四个核心交付可解析且冻结输入哈希正确", 4, 4)]
    scores = {"R001": 1.0}
    for key in [f"R{index:03d}" for index in range(2, 13)]:
        score, failures, earned, total = groups[key]
        evidence = "全部任务专属原子检查通过" if not failures else "未通过：" + "；".join(str(item) for item in failures)
        results.append(_criterion(key, score, evidence, earned, total))
        scores[key] = score

    penalties = {
        "R013": _identity_penalty(delivered_ids, [str(row["order_id"]) for row in rows]),
        "R014": _cross_penalty(artifact, attachments, word_text, mail_body, calendar_description, facts),
        "R015": _formula_penalty(formula_count, formula_total),
        "R016": _boundary_penalty("\n".join((word_text, mail_body, calendar_text))),
        "R017": _safe_penalty(artifact, "\n".join((word_text, mail_body, calendar_text))),
    }
    for key, (severity, labels, total) in penalties.items():
        results.append(_criterion(key, severity, "未触发任务专属严重边界" if not labels else "；".join(labels), severity * total, total))
        scores[key] = severity

    positive_total = sum(POSITIVE_WEIGHTS.values())
    positive = sum(POSITIVE_WEIGHTS[key] * scores[key] for key in POSITIVE_WEIGHTS)
    penalty = sum(abs(PENALTY_WEIGHTS[key]) * scores[key] for key in PENALTY_WEIGHTS)
    reward = round(max(0.0, min(1.0, (positive - penalty) / positive_total)), 6)
    workbook.close()
    return {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes, "criterion_results": results, "reward": reward,
        "passed": reward >= 0.8, "hurdle_failures": [],
        "penalties": [{"id": key, "amount": round(abs(PENALTY_WEIGHTS[key]) * scores[key] / positive_total, 6)}
                      for key in PENALTY_WEIGHTS if scores[key] > 0],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    parser.add_argument("--scoring-version", choices=("v1", "v2"), default="v2")
    args = parser.parse_args()
    try:
        result = V1.evaluate(args.artifact_dir) if args.scoring_version == "v1" else evaluate(args.artifact_dir)
    except Exception as exc:  # noqa: BLE001
        result = _fatal([f"Judge 异常：{type(exc).__name__}: {exc}"], [])
    output = json.dumps(result, ensure_ascii=False, indent=2)
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(output + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{float(result['reward']):.6f}\n", encoding="utf-8")
    print(output)


if __name__ == "__main__":
    main()
