#!/usr/bin/env python3
"""跨系统财务审计任务的来源重算、Office 结构和跨文件一致性 Judge。"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import posixpath
import re
import shutil
import subprocess
import tempfile
import zipfile
from datetime import date, datetime
from email import policy
from email.parser import BytesParser
from pathlib import Path

from docx import Document
from lxml import etree
from openpyxl import load_workbook


TASK_ID = "cowork-evo/cowork-bench-financial-audit-reconciliation-package"
RUBRIC_VERSION = "cowork-bench-financial-audit-reconciliation-v0.1"
JUDGE_VERSION = "cowork-bench-financial-audit-reconciliation-deterministic-v0.1"
XLSX_NAME = "Audit_Reconciliation_Report.xlsx"
DOCX_NAME = "Audit_Report.docx"
EML_NAME = "Financial_Audit_Reconciliation.eml"
ICS_NAME = "Audit_Reconciliation_Review.ics"
REQUIRED_FILES = [XLSX_NAME, DOCX_NAME, EML_NAME, ICS_NAME]
SOURCE_HASHES = {
    "authoritative_ledger.csv": "6d141be371f2b2d8942e493b1f59b4a631ccd0eb482bb2ced6bcc3dd4989cb00",
    "q1_manual_tracking.xlsx": "d980f32983abca1921f64282f9fb6d347e2080dcc7ac9ee721f7a8c43719d95e",
    "system_mapping_guide.pdf": "457fd78134f2ea0877129436a13bb1cbce79270bf9f991d57a3139c03f61fddf",
}
ALLOWED_STATUS = {"completed", "pending"}
RECON_HEADERS = [
    "Order ID", "Manual Date", "Ledger Date", "Date Delta Days", "Manual Amount",
    "Ledger Amount", "Amount Delta", "Manual Status", "Ledger Status", "Presence Check",
    "Date Check", "Amount Check", "Status Check", "Match Status", "Discrepancy Reason",
]
EXCEPTION_HEADERS = ["Order ID", "Exception Type", "Manual Evidence", "Ledger Evidence", "Impact", "Required Follow-up"]
SUMMARY_METRICS = [
    "Total Manual Orders", "Matched", "Discrepancies", "Manual Total Amount",
    "Authoritative Amount Represented", "Missing Authoritative Amount", "Absolute Amount Mismatch",
    "Date Mismatches", "Amount Mismatches", "Missing Records", "Control Conclusion",
]
WORD_HEADINGS = [
    "Executive Summary", "Scope and Methodology", "Reconciliation Results",
    "Discrepancies and Required Actions", "Evidence Boundary",
]
SUSPICIOUS_MARKERS = [
    "solution/golden", "judge.py", "expectations.json", "judge-result", "/root/", "/home/", "c:\\",
]

SPECS = {
    "R001": ("完整交付物与冻结输入", 6, "hurdle", 0.0),
    "R002": ("Reconciliation 来源字段与公式", 18, "hurdle", 0.3),
    "R003": ("Summary 公式与汇总真值", 10, "hurdle", 0.4),
    "R004": ("Exception Register 事实与行动", 10, "hurdle", 0.45),
    "R005": ("Word 范围、方法与总体结果", 8, "hurdle", 0.5),
    "R006": ("Word 异常事实与行动", 10, "hurdle", 0.5),
    "R007": ("EML 正文与字节一致附件", 12, "hurdle", 0.45),
    "R008": ("ICS 会议邀请", 8, "hurdle", 0.55),
    "R009": ("四文件事实一致", 10, "hurdle", 0.49),
    "R010": ("原生 Office 结构", 5, "hurdle", 0.65),
    "R011": ("LibreOffice 渲染与可用性", 4, "positive", None),
    "R012": ("证据边界与无来源断言", 5, "hurdle", 0.7),
    "R013": ("安全、干净交付", 3, "positive", None),
    "R014": ("无来源真值冲突", -25, "penalty", None),
    "R015": ("无跨文件冲突或陈旧附件", -25, "penalty", None),
    "R016": ("无缺失、损坏、静态化或截图化", -25, "penalty", None),
    "R017": ("无不支持的业务结论", -15, "penalty", None),
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def normalize(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\u00a0", " ").replace("\v", " ")).strip()


def lower(value: object) -> str:
    return normalize(value).lower()


def close_number(value: object, expected: float, tolerance: float = 0.005) -> bool:
    try:
        if value is None or isinstance(value, bool):
            return False
        return abs(float(value) - expected) <= tolerance
    except (TypeError, ValueError):
        return False


def iso_date(value: object) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = normalize(value)
    try:
        return date.fromisoformat(text[:10]).isoformat()
    except ValueError:
        return None


def source_path(name: str) -> Path | None:
    task_root = Path(__file__).resolve().parents[1]
    for candidate in (Path("/app/input_files") / name, task_root / "data" / "input_files" / name):
        if candidate.is_file():
            return candidate
    return None


def source_integrity() -> tuple[bool, dict[str, str]]:
    errors: dict[str, str] = {}
    for name, expected in SOURCE_HASHES.items():
        path = source_path(name)
        if path is None:
            errors[name] = "missing"
        elif sha256(path) != expected:
            errors[name] = f"sha256={sha256(path)}"
    return not errors, errors


def compute_truth() -> tuple[list[dict[str, object]], dict[str, object]]:
    ledger_path = source_path("authoritative_ledger.csv")
    manual_path = source_path("q1_manual_tracking.xlsx")
    if ledger_path is None or manual_path is None:
        raise FileNotFoundError("缺少来源 CSV/XLSX")
    with ledger_path.open(encoding="utf-8", newline="") as handle:
        ledger = {row["Order ID"].strip(): row for row in csv.DictReader(handle)}
    workbook = load_workbook(manual_path, data_only=True)
    sheet = workbook["Manual Tracking"]
    headers = [normalize(cell.value) for cell in sheet[1]]
    rows: list[dict[str, object]] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        manual = dict(zip(headers, values, strict=True))
        order_id = normalize(manual["Order ID"])
        source = ledger.get(order_id)
        manual_date = date.fromisoformat(normalize(manual["Date"]))
        manual_amount = round(float(manual["Amount"]), 2)
        manual_status = normalize(manual["Status"])
        ledger_date = date.fromisoformat(source["Date"]) if source else None
        ledger_amount = round(float(source["Amount"]), 2) if source else None
        ledger_status = source["Status"].strip() if source else None
        date_delta = abs((manual_date - ledger_date).days) if ledger_date else None
        amount_delta = round(manual_amount - ledger_amount, 2) if ledger_amount is not None else None
        reasons: list[str] = []
        if source is None:
            reasons.append("missing_in_authoritative_ledger")
        else:
            if date_delta is None or date_delta > 1:
                reasons.append("date_mismatch")
            if amount_delta is None or abs(amount_delta) >= 0.005:
                reasons.append("amount_mismatch")
            if manual_status.lower() not in ALLOWED_STATUS or str(ledger_status).lower() not in ALLOWED_STATUS:
                reasons.append("invalid_status")
        rows.append(
            {
                "order_id": order_id,
                "manual_date": manual_date.isoformat(),
                "ledger_date": ledger_date.isoformat() if ledger_date else None,
                "date_delta": date_delta,
                "manual_amount": manual_amount,
                "ledger_amount": ledger_amount,
                "amount_delta": amount_delta,
                "manual_status": manual_status,
                "ledger_status": ledger_status,
                "presence_check": "pass" if source else "fail",
                "date_check": "pass" if source and date_delta is not None and date_delta <= 1 else "fail",
                "amount_check": "pass" if source and amount_delta is not None and abs(amount_delta) < 0.005 else "fail",
                "status_check": "pass" if source and manual_status.lower() in ALLOWED_STATUS and str(ledger_status).lower() in ALLOWED_STATUS else "fail",
                "match_status": "matched" if not reasons else "discrepancy",
                "reason": ";".join(reasons) if reasons else "none",
            }
        )
    rows.sort(key=lambda row: str(row["order_id"]))
    discrepancies = [row for row in rows if row["match_status"] == "discrepancy"]
    facts = {
        "total": len(rows),
        "matched": sum(row["match_status"] == "matched" for row in rows),
        "discrepancies": len(discrepancies),
        "manual_total": round(sum(float(row["manual_amount"]) for row in rows), 2),
        "ledger_total": round(sum(float(row["ledger_amount"]) for row in rows if row["ledger_amount"] is not None), 2),
        "missing_amount": round(sum(float(row["manual_amount"]) for row in rows if row["presence_check"] == "fail"), 2),
        "absolute_amount_mismatch": round(sum(abs(float(row["amount_delta"])) for row in rows if "amount_mismatch" in str(row["reason"])), 2),
        "date_mismatches": sum("date_mismatch" in str(row["reason"]) for row in rows),
        "amount_mismatches": sum("amount_mismatch" in str(row["reason"]) for row in rows),
        "missing_records": sum("missing_in_authoritative_ledger" in str(row["reason"]) for row in rows),
        "discrepancy_rows": discrepancies,
        "discrepancy_ids": [str(row["order_id"]) for row in discrepancies],
    }
    return rows, facts


def ooxml_package(path: Path) -> tuple[bool, dict[str, object]]:
    evidence: dict[str, object] = {"broken": None, "missing_targets": [], "macros": [], "external": [], "media": []}
    try:
        with zipfile.ZipFile(path) as archive:
            evidence["broken"] = archive.testzip()
            names = set(archive.namelist())
            if "[Content_Types].xml" not in names:
                evidence["missing_targets"].append("[Content_Types].xml")
            evidence["macros"] = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
            evidence["media"] = sorted(name for name in names if "/media/" in name and not name.endswith("/"))
            for rel_name in sorted(name for name in names if name.endswith(".rels")):
                root = etree.fromstring(archive.read(rel_name))
                if rel_name == "_rels/.rels":
                    owner = ""
                else:
                    before, after = rel_name.split("/_rels/", 1)
                    owner = posixpath.join(before, after[:-5])
                base = posixpath.dirname(owner)
                for relation in root:
                    target = relation.attrib.get("Target", "")
                    if relation.attrib.get("TargetMode") == "External":
                        evidence["external"].append(target)
                        continue
                    resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                    if resolved not in names:
                        evidence["missing_targets"].append(f"{rel_name}->{resolved}")
        ok = not evidence["broken"] and not evidence["missing_targets"] and not evidence["macros"]
        return ok, evidence
    except Exception as error:  # noqa: BLE001 - malformed submissions must score, not crash.
        evidence["exception"] = str(error)
        return False, evidence


def recalculate_workbook(path: Path):
    formulas = load_workbook(path, data_only=False)
    with tempfile.TemporaryDirectory(prefix="audit-judge-recalc-") as tmp:
        root = Path(tmp)
        input_root = root / "input"
        output_root = root / "output"
        profile = root / "profile"
        input_root.mkdir()
        output_root.mkdir()
        profile.mkdir()
        source = input_root / path.name
        shutil.copy2(path, source)
        result = subprocess.run(
            ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "xlsx", "--outdir", str(output_root), str(source)],
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            timeout=300,
        )
        converted = output_root / path.name
        if result.returncode != 0 or not converted.is_file():
            raise RuntimeError(result.stdout[-1000:])
        values = load_workbook(converted, data_only=True)
    return formulas, values


def formula_cell(cell) -> bool:
    return cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("="))


def workbook_checks(path: Path, expected_rows: list[dict[str, object]], facts: dict[str, object]) -> tuple[dict[str, bool], dict[str, object]]:
    checks = {"reconciliation": False, "summary": False, "exceptions": False, "native": False}
    evidence: dict[str, object] = {}
    try:
        formulas, values = recalculate_workbook(path)
        evidence["sheetnames"] = formulas.sheetnames
        if formulas.sheetnames != ["Reconciliation", "Exception Register", "Summary"]:
            return checks, evidence

        f_recon = formulas["Reconciliation"]
        v_recon = values["Reconciliation"]
        headers = [normalize(f_recon.cell(1, column).value) for column in range(1, 16)]
        delivered_ids = [normalize(v_recon.cell(row, 1).value) for row in range(2, v_recon.max_row + 1) if normalize(v_recon.cell(row, 1).value)]
        formula_positions = [(row, column) for row in range(2, 17) for column in (4, 7, 10, 11, 12, 13, 14, 15)]
        formulas_ok = all(formula_cell(f_recon.cell(row, column)) for row, column in formula_positions)
        row_errors: list[str] = []
        if headers != RECON_HEADERS:
            row_errors.append(f"headers={headers}")
        if delivered_ids != [str(row["order_id"]) for row in expected_rows]:
            row_errors.append(f"ids={delivered_ids}")
        for row_number, expected in enumerate(expected_rows, start=2):
            actual = [v_recon.cell(row_number, column).value for column in range(1, 16)]
            conditions = [
                normalize(actual[0]) == expected["order_id"],
                iso_date(actual[1]) == expected["manual_date"],
                iso_date(actual[2]) == expected["ledger_date"],
                (actual[3] in (None, "") if expected["date_delta"] is None else close_number(actual[3], float(expected["date_delta"]))),
                close_number(actual[4], float(expected["manual_amount"])),
                (actual[5] in (None, "") if expected["ledger_amount"] is None else close_number(actual[5], float(expected["ledger_amount"]))),
                (actual[6] in (None, "") if expected["amount_delta"] is None else close_number(actual[6], float(expected["amount_delta"]))),
                lower(actual[7]) == lower(expected["manual_status"]),
                (actual[8] in (None, "") if expected["ledger_status"] is None else lower(actual[8]) == lower(expected["ledger_status"])),
                lower(actual[9]) == expected["presence_check"],
                lower(actual[10]) == expected["date_check"],
                lower(actual[11]) == expected["amount_check"],
                lower(actual[12]) == expected["status_check"],
                lower(actual[13]) == expected["match_status"],
                lower(actual[14]) == expected["reason"],
            ]
            if not all(conditions):
                row_errors.append(f"{expected['order_id']}={actual}")
        checks["reconciliation"] = formulas_ok and not row_errors
        evidence["reconciliation_formula_count"] = sum(formula_cell(f_recon.cell(row, column)) for row, column in formula_positions)
        evidence["reconciliation_errors"] = row_errors[:12]

        f_summary = formulas["Summary"]
        v_summary = values["Summary"]
        summary_headers = [normalize(v_summary.cell(1, column).value) for column in (1, 2)]
        metrics = [normalize(v_summary.cell(row, 1).value) for row in range(2, 13)]
        summary_formulas = all(formula_cell(f_summary.cell(row, 2)) for row in range(2, 13))
        expected_values: list[object] = [
            facts["total"], facts["matched"], facts["discrepancies"], facts["manual_total"], facts["ledger_total"],
            facts["missing_amount"], facts["absolute_amount_mismatch"], facts["date_mismatches"], facts["amount_mismatches"],
            facts["missing_records"], "Needs remediation",
        ]
        actual_values = [v_summary.cell(row, 2).value for row in range(2, 13)]
        summary_values_ok = all(
            (close_number(actual, float(expected)) if isinstance(expected, (int, float)) else lower(actual) == lower(expected))
            for actual, expected in zip(actual_values, expected_values, strict=True)
        )
        checks["summary"] = summary_headers == ["Metric", "Value"] and metrics == SUMMARY_METRICS and summary_formulas and summary_values_ok
        evidence["summary"] = {"headers": summary_headers, "metrics": metrics, "values": actual_values, "formula_count": sum(formula_cell(f_summary.cell(row, 2)) for row in range(2, 13))}

        v_exceptions = values["Exception Register"]
        exception_headers = [normalize(v_exceptions.cell(1, column).value) for column in range(1, 7)]
        exception_rows = {
            normalize(v_exceptions.cell(row, 1).value): [normalize(v_exceptions.cell(row, column).value) for column in range(1, 7)]
            for row in range(2, v_exceptions.max_row + 1)
            if normalize(v_exceptions.cell(row, 1).value)
        }
        exception_errors: list[str] = []
        if exception_headers != EXCEPTION_HEADERS:
            exception_errors.append(f"headers={exception_headers}")
        if set(exception_rows) != set(facts["discrepancy_ids"]):
            exception_errors.append(f"ids={sorted(exception_rows)}")
        for expected in facts["discrepancy_rows"]:
            delivered = exception_rows.get(str(expected["order_id"]), [])
            if not delivered:
                continue
            row_text = lower(" | ".join(delivered))
            reason = str(expected["reason"])
            if reason == "amount_mismatch":
                type_ok = "amount" in row_text or "金额" in row_text
                numbers_ok = all(f"{value:.2f}" in row_text for value in (float(expected["manual_amount"]), float(expected["ledger_amount"]), abs(float(expected["amount_delta"]))))
            elif reason == "date_mismatch":
                type_ok = "date" in row_text or "日期" in row_text
                numbers_ok = expected["manual_date"] in row_text and expected["ledger_date"] in row_text and str(expected["date_delta"]) in row_text
            else:
                type_ok = ("missing" in row_text or "缺失" in row_text or "无对应" in row_text) and ("authoritative" in row_text or "权威" in row_text)
                numbers_ok = f"{float(expected['manual_amount']):.2f}" in row_text
            action = lower(delivered[5])
            action_ok = len(action) >= 20 and any(token in action for token in ("verify", "validate", "investigate", "correct", "核验", "调查", "更正"))
            if not (type_ok and numbers_ok and action_ok):
                exception_errors.append(f"{expected['order_id']}={delivered}")
        checks["exceptions"] = not exception_errors
        evidence["exception_errors"] = exception_errors

        native_tables = sum(len(sheet.tables) for sheet in formulas.worksheets)
        freeze_ok = all(formulas[name].freeze_panes == "A2" for name in formulas.sheetnames)
        filters_ok = all(bool(formulas[name].auto_filter.ref) for name in formulas.sheetnames)
        conditional_count = sum(len(sheet.conditional_formatting) for sheet in formulas.worksheets)
        checks["native"] = native_tables >= 3 and freeze_ok and filters_ok and conditional_count >= 2 and formulas_ok and summary_formulas
        evidence["native"] = {"tables": native_tables, "freeze": freeze_ok, "filters": filters_ok, "conditional_formats": conditional_count}
    except Exception as error:  # noqa: BLE001
        evidence["exception"] = str(error)
    return checks, evidence


def document_text(document: Document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def has_money(text: str, value: float) -> bool:
    variants = {
        f"{value:,.2f}", f"{value:.2f}", f"{value:,.2f}".replace(",", " "),
        f"{value:,.2f}".replace(",", "\u00a0"),
    }
    return any(variant in text for variant in variants)


def count_with_label(text: str, label: str, count: int) -> bool:
    """要求数量和语义标签相邻，避免把 WC-011 中的 11 当成 matched 数。"""
    normalized = lower(text)
    if label == "orders":
        patterns = [
            rf"\b{count}\s+(?:manual[- ]tracking\s+|manual\s+)?orders\b",
            rf"\bmanual orders\b[^0-9]{{0,24}}\b{count}\b",
        ]
    elif label == "matched":
        patterns = [
            rf"\b{count}\s+(?:orders\s+)?matched\b",
            rf"\bmatched\b[^0-9]{{0,24}}\b{count}\b",
        ]
    else:
        patterns = [
            rf"\b{count}\s+(?:orders\s+)?discrepanc(?:y|ies)\b",
            rf"\bdiscrepanc(?:y|ies)\b[^0-9]{{0,24}}\b{count}\b",
        ]
    return any(re.search(pattern, normalized) for pattern in patterns)


def reason_present(text: str, reason: str) -> bool:
    lowered = lower(text)
    if reason == "amount_mismatch":
        return "amount_mismatch" in lowered or "amount mismatch" in lowered or "金额不一致" in lowered or "金额差异" in lowered
    if reason == "date_mismatch":
        return "date_mismatch" in lowered or "date mismatch" in lowered or "日期不一致" in lowered or "日期差异" in lowered
    if reason == "missing_in_authoritative_ledger":
        return "missing_in_authoritative_ledger" in lowered or "missing in authoritative" in lowered or "权威台账缺失" in lowered or "权威台账中无" in lowered
    return "invalid_status" in lowered or "invalid status" in lowered or "状态无效" in lowered


def word_checks(path: Path, facts: dict[str, object]) -> tuple[dict[str, bool], dict[str, object], str]:
    checks = {"summary_method": False, "exceptions": False, "native": False}
    evidence: dict[str, object] = {}
    text = ""
    try:
        document = Document(path)
        text = document_text(document)
        lowered = lower(text)
        headings = [normalize(paragraph.text) for paragraph in document.paragraphs if paragraph.style.name == "Heading 1"]
        method_tokens = ["authoritative", "1 day", "zero tolerance", "case-insensitive", "completed", "pending"]
        summary_ok = (
            all(heading in headings for heading in WORD_HEADINGS)
            and count_with_label(text, "orders", int(facts["total"]))
            and count_with_label(text, "matched", int(facts["matched"]))
            and count_with_label(text, "discrepancies", int(facts["discrepancies"]))
            and has_money(text, float(facts["manual_total"]))
            and all(token in lowered for token in method_tokens)
        )
        checks["summary_method"] = summary_ok
        exception_errors = []
        for row in facts["discrepancy_rows"]:
            order_id = str(row["order_id"])
            if order_id not in text or not reason_present(text, str(row["reason"])):
                exception_errors.append(order_id)
                continue
            if row["reason"] == "amount_mismatch":
                if not all(has_money(text, float(value)) for value in (row["manual_amount"], row["ledger_amount"], abs(float(row["amount_delta"])))):
                    exception_errors.append(order_id)
            elif row["reason"] == "date_mismatch":
                if row["manual_date"] not in text or row["ledger_date"] not in text or str(row["date_delta"]) not in text:
                    exception_errors.append(order_id)
            elif not has_money(text, float(row["manual_amount"])):
                exception_errors.append(order_id)
        action_tokens = ("verify", "validate", "investigate", "correct", "核验", "调查", "更正")
        checks["exceptions"] = not exception_errors and any(token in lowered for token in action_tokens)
        exception_tables = [table for table in document.tables if len(table.columns) >= 6 and any("order" in lower(cell.text) for cell in table.rows[0].cells)]
        checks["native"] = len(document.tables) >= 3 and len(exception_tables) >= 1 and len(headings) >= 5
        evidence = {"headings": headings, "tables": len(document.tables), "exception_tables": len(exception_tables), "exception_errors": exception_errors, "text_length": len(text)}
    except Exception as error:  # noqa: BLE001
        evidence["exception"] = str(error)
    return checks, evidence, text


def email_checks(path: Path, artifact_dir: Path, facts: dict[str, object]) -> tuple[bool, dict[str, object], str]:
    evidence: dict[str, object] = {}
    body = ""
    try:
        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        recipients = [address.addr_spec.lower() for address in message["to"].addresses] if message.get("to") else []
        subject = normalize(message.get("subject"))
        for part in message.walk():
            if part.get_content_type() == "text/plain" and part.get_content_disposition() != "attachment":
                body += part.get_content()
        attachments = {}
        attachment_types = {}
        for part in message.iter_attachments():
            name = part.get_filename()
            if name:
                attachments[name] = part.get_payload(decode=True) or b""
                attachment_types[name] = part.get_content_type()
        attachment_ok = set(attachments) == {XLSX_NAME, DOCX_NAME}
        for name in (XLSX_NAME, DOCX_NAME):
            delivered = artifact_dir / name
            attachment_ok = attachment_ok and delivered.is_file() and attachments.get(name) == delivered.read_bytes()
        body_ok = (
            count_with_label(body, "orders", int(facts["total"]))
            and count_with_label(body, "matched", int(facts["matched"]))
            and count_with_label(body, "discrepancies", int(facts["discrepancies"]))
            and has_money(body, float(facts["manual_total"]))
            and all(order_id in body for order_id in facts["discrepancy_ids"])
            and all(reason_present(body, str(row["reason"])) for row in facts["discrepancy_rows"])
        )
        valid = (
            recipients == ["finance@company.com"]
            and subject == "Financial Audit Reconciliation - Q1 2024"
            and lower(message.get("X-Draft-Status")) == "draft"
            and body_ok
            and attachment_ok
        )
        evidence = {"recipients": recipients, "subject": subject, "draft": message.get("X-Draft-Status"), "attachments": {name: len(data) for name, data in attachments.items()}, "attachment_types": attachment_types, "body_ok": body_ok}
        return valid, evidence, body
    except Exception as error:  # noqa: BLE001
        evidence["exception"] = str(error)
        return False, evidence, body


def unfold_ics(text: str) -> list[str]:
    lines: list[str] = []
    for raw in text.replace("\r\n", "\n").split("\n"):
        if raw.startswith((" ", "\t")) and lines:
            lines[-1] += raw[1:]
        elif raw:
            lines.append(raw)
    return lines


def calendar_checks(path: Path, facts: dict[str, object]) -> tuple[bool, dict[str, object], str]:
    evidence: dict[str, object] = {}
    text = ""
    try:
        text = path.read_text(encoding="utf-8")
        lines = unfold_ics(text)
        properties: dict[str, list[str]] = {}
        for line in lines:
            if ":" not in line:
                continue
            key, value = line.split(":", 1)
            properties.setdefault(key, []).append(value)
        description = " ".join(value for key, values in properties.items() if key == "DESCRIPTION" for value in values)
        attachments = [value for key, values in properties.items() if key.startswith("ATTACH") for value in values]
        attendee_values = [value.lower() for key, values in properties.items() if key.startswith("ATTENDEE") for value in values]
        valid = all(
            [
                "BEGIN:VCALENDAR" in lines,
                "BEGIN:VEVENT" in lines,
                properties.get("DTSTART;TZID=Asia/Shanghai") == ["20240405T100000"],
                properties.get("DTEND;TZID=Asia/Shanghai") == ["20240405T104500"],
                properties.get("SUMMARY") == ["Audit Reconciliation Review"],
                properties.get("STATUS") == ["TENTATIVE"],
                attendee_values == ["mailto:finance@company.com"],
                set(attachments) == {XLSX_NAME, DOCX_NAME},
                count_with_label(description, "orders", int(facts["total"])),
                count_with_label(description, "matched", int(facts["matched"])),
                count_with_label(description, "discrepancies", int(facts["discrepancies"])),
                all(order_id in description for order_id in facts["discrepancy_ids"]),
                all(reason_present(description, str(row["reason"])) for row in facts["discrepancy_rows"]),
                not any(marker in lower(text) for marker in ("file://", "/app/", "/home/", "c:\\")),
            ]
        )
        evidence = {"start": properties.get("DTSTART;TZID=Asia/Shanghai"), "end": properties.get("DTEND;TZID=Asia/Shanghai"), "summary": properties.get("SUMMARY"), "status": properties.get("STATUS"), "attendees": attendee_values, "attachments": attachments, "description": description}
        return valid, evidence, text
    except Exception as error:  # noqa: BLE001
        evidence["exception"] = str(error)
        return False, evidence, text


def render_office(path: Path) -> tuple[int, str, str]:
    with tempfile.TemporaryDirectory(prefix="audit-judge-render-") as tmp:
        root = Path(tmp)
        output = root / "output"
        profile = root / "profile"
        output.mkdir()
        profile.mkdir()
        result = subprocess.run(
            ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf", "--outdir", str(output), str(path)],
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            timeout=300,
        )
        pdf = output / f"{path.stem}.pdf"
        if result.returncode != 0 or not pdf.is_file():
            return 0, "", result.stdout[-1000:]
        info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=60).stdout
        match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
        pages = int(match.group(1)) if match else 0
        text = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=60).stdout
        return pages, text, result.stdout[-1000:]


def unsupported_claims(text: str) -> list[str]:
    normalized = lower(text)
    patterns = {
        "fraud assertion": r"\b(?:(?:is|was)\s+)?(?:confirmed|classified)(?:\s+as|\s+to\s+be)?\s+(?:a\s+)?fraud(?:ulent)?\b|\b(?:is|was)\s+(?:a\s+)?fraud(?:ulent)?\b",
        "cancel assertion": r"\b(?:is|was|has been)\s+cancel(?:led|ed)\b",
        "remediation complete": r"\bremediation\s+(?:is|was|has been)\s+complete(?:d)?\b",
        "clean audit opinion": r"financial statements\s+(?:are|were)\s+free of material misstatement",
        "unsupported quantified impact": r"(?:revenue impact|penalt(?:y|ies)|customer loss|roi|responsibility (?:share|allocation))\s*[:=]\s*(?:rub|\$|\d)",
    }
    findings: list[str] = []
    for name, pattern in patterns.items():
        for match in re.finditer(pattern, normalized):
            context = normalized[max(0, match.start() - 90) : min(len(normalized), match.end() + 40)]
            if any(negation in context for negation in ("does not", "do not", "not support", "cannot support", "whether")):
                continue
            findings.append(name)
            break
    return findings


def item(criterion_id: str, passed: bool, evidence: object) -> dict[str, object]:
    description, weight, kind, cap = SPECS[criterion_id]
    raw: dict[str, object] = {"description": description, "weight": weight, "type": kind, "passed": bool(passed)}
    if cap is not None:
        raw["score_cap_on_failure"] = cap
    return {
        "criterion_id": criterion_id,
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
        "raw": raw,
    }


def assemble(artifact_dir: Path, criteria: list[dict[str, object]]) -> dict[str, object]:
    positive_total = sum(float(result["raw"]["weight"]) for result in criteria if float(result["raw"]["weight"]) > 0)
    earned = sum(float(result["raw"]["weight"]) * float(result["score"]) for result in criteria if float(result["raw"]["weight"]) > 0)
    penalty = sum(abs(float(result["raw"]["weight"])) for result in criteria if float(result["raw"]["weight"]) < 0 and not result["raw"]["passed"])
    reward = max(0.0, min(1.0, (earned - penalty) / positive_total if positive_total else 0.0))
    hurdle_failures: list[str] = []
    for result in criteria:
        if result["raw"]["type"] == "hurdle" and not result["raw"]["passed"]:
            hurdle_failures.append(str(result["criterion_id"]))
            reward = min(reward, float(result["raw"].get("score_cap_on_failure", 1.0)))
    hashes = [sha256(artifact_dir / name) for name in REQUIRED_FILES if (artifact_dir / name).is_file()]
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes,
        "criterion_results": criteria,
        "reward": round(reward, 6),
        "passed": reward >= 0.8,
        "hurdle_failures": hurdle_failures,
    }


def evaluate(artifact_dir: Path) -> dict[str, object]:
    source_ok, source_errors = source_integrity()
    expected_rows, facts = compute_truth()
    delivered = sorted(str(path.relative_to(artifact_dir)) for path in artifact_dir.rglob("*") if path.is_file()) if artifact_dir.is_dir() else []
    paths = {name: artifact_dir / name for name in REQUIRED_FILES}
    required_ok = delivered == sorted(REQUIRED_FILES)
    package_results = {}
    package_ok = True
    for name in (XLSX_NAME, DOCX_NAME):
        if paths[name].is_file():
            ok, evidence = ooxml_package(paths[name])
        else:
            ok, evidence = False, {"missing": name}
        package_results[name] = evidence
        package_ok = package_ok and ok
    r1 = source_ok and required_ok and package_ok
    criteria = [item("R001", r1, {"source_errors": source_errors, "delivered": delivered, "packages": package_results})]

    workbook_state = {"reconciliation": False, "summary": False, "exceptions": False, "native": False}
    workbook_evidence: dict[str, object] = {"missing": XLSX_NAME}
    if paths[XLSX_NAME].is_file():
        workbook_state, workbook_evidence = workbook_checks(paths[XLSX_NAME], expected_rows, facts)
    criteria.append(item("R002", workbook_state["reconciliation"], workbook_evidence))
    criteria.append(item("R003", workbook_state["summary"], workbook_evidence.get("summary", workbook_evidence)))
    criteria.append(item("R004", workbook_state["exceptions"], workbook_evidence.get("exception_errors", workbook_evidence)))

    word_state = {"summary_method": False, "exceptions": False, "native": False}
    word_evidence: dict[str, object] = {"missing": DOCX_NAME}
    word_text = ""
    if paths[DOCX_NAME].is_file():
        word_state, word_evidence, word_text = word_checks(paths[DOCX_NAME], facts)
    criteria.append(item("R005", word_state["summary_method"], word_evidence))
    criteria.append(item("R006", word_state["exceptions"], word_evidence.get("exception_errors", word_evidence)))

    email_ok, email_evidence, email_body = (False, {"missing": EML_NAME}, "")
    if paths[EML_NAME].is_file():
        email_ok, email_evidence, email_body = email_checks(paths[EML_NAME], artifact_dir, facts)
    criteria.append(item("R007", email_ok, email_evidence))

    calendar_ok, calendar_evidence, calendar_text = (False, {"missing": ICS_NAME}, "")
    if paths[ICS_NAME].is_file():
        calendar_ok, calendar_evidence, calendar_text = calendar_checks(paths[ICS_NAME], facts)
    criteria.append(item("R008", calendar_ok, calendar_evidence))

    cross_ok = all([workbook_state["reconciliation"], workbook_state["summary"], workbook_state["exceptions"], word_state["summary_method"], word_state["exceptions"], email_ok, calendar_ok])
    criteria.append(item("R009", cross_ok, {"workbook": workbook_state, "word": word_state, "email": email_ok, "calendar": calendar_ok}))

    media = package_results.get(XLSX_NAME, {}).get("media", []) + package_results.get(DOCX_NAME, {}).get("media", [])
    native_ok = workbook_state["native"] and word_state["native"] and not media
    criteria.append(item("R010", native_ok, {"workbook": workbook_evidence.get("native"), "word": word_evidence, "media": media}))

    render_evidence = {}
    render_ok = False
    if paths[XLSX_NAME].is_file() and paths[DOCX_NAME].is_file():
        try:
            xlsx_pages, xlsx_text, xlsx_log = render_office(paths[XLSX_NAME])
            docx_pages, docx_text, docx_log = render_office(paths[DOCX_NAME])
            xlsx_tokens = ["Discrepancy Reason", "Required Follow-up", "Control Conclusion", "WC-013"]
            docx_tokens = WORD_HEADINGS + facts["discrepancy_ids"]
            render_ok = 3 <= xlsx_pages <= 6 and 2 <= docx_pages <= 5 and all(token in xlsx_text for token in xlsx_tokens) and all(token in docx_text for token in docx_tokens)
            render_evidence = {"xlsx_pages": xlsx_pages, "docx_pages": docx_pages, "xlsx_missing": [token for token in xlsx_tokens if token not in xlsx_text], "docx_missing": [token for token in docx_tokens if token not in docx_text], "logs": [xlsx_log, docx_log]}
        except Exception as error:  # noqa: BLE001
            render_evidence = {"exception": str(error)}
    criteria.append(item("R011", render_ok, render_evidence))

    combined_text = "\n".join([word_text, email_body, calendar_text])
    unsupported = unsupported_claims(combined_text)
    boundary_tokens = ["does not", "fraud", "cancellation", "material misstatement", "verification"]
    boundary_ok = not unsupported and all(token in lower(word_text) for token in boundary_tokens)
    criteria.append(item("R012", boundary_ok, {"unsupported": unsupported, "boundary_missing": [token for token in boundary_tokens if token not in lower(word_text)]}))

    external = package_results.get(XLSX_NAME, {}).get("external", []) + package_results.get(DOCX_NAME, {}).get("external", [])
    suspicious = [marker for marker in SUSPICIOUS_MARKERS if marker in lower(combined_text)]
    clean_ok = required_ok and package_ok and not external and not suspicious
    criteria.append(item("R013", clean_ok, {"delivered": delivered, "external": external, "suspicious": suspicious}))

    source_truth_ok = all([workbook_state["reconciliation"], workbook_state["summary"], workbook_state["exceptions"], word_state["summary_method"], word_state["exceptions"]])
    criteria.append(item("R014", source_truth_ok, {"source_truth_ok": source_truth_ok}))
    criteria.append(item("R015", cross_ok and email_ok, {"cross_ok": cross_ok, "email_attachments_ok": email_ok}))
    no_flattening = package_ok and workbook_state["native"] and word_state["native"] and not media
    criteria.append(item("R016", required_ok and no_flattening, {"required_ok": required_ok, "native": no_flattening, "media": media}))
    criteria.append(item("R017", boundary_ok, {"unsupported": unsupported}))
    return assemble(artifact_dir, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    result = evaluate(args.artifact_dir)
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{result['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
