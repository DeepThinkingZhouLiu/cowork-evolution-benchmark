#!/usr/bin/env python3
"""Q1 财务审计包的任务专属 v2 致命与分层局部反例。"""

from __future__ import annotations

import importlib.util
import json
import shutil
import sys
import tempfile
from email import encoders, policy
from email.parser import BytesParser
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
REQUIRED = (
    "Audit_Reconciliation_Report.xlsx",
    "Audit_Report.docx",
    "Financial_Audit_Reconciliation.eml",
    "Audit_Reconciliation_Review.ics",
)
XLSX, DOCX, EML, ICS = REQUIRED


def _load_judge():
    path = Path(__file__).with_name("judge.py")
    spec = importlib.util.spec_from_file_location("financial_audit_v2_negative_judge", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


JUDGE = _load_judge()


def copy_golden(target: Path) -> None:
    target.mkdir(parents=True, exist_ok=True)
    for name in REQUIRED:
        shutil.copy2(GOLDEN / name, target / name)


def sync_email_attachments(root: Path) -> None:
    path = root / EML
    message = BytesParser(policy=policy.default.clone(linesep="\r\n")).parsebytes(path.read_bytes())
    for part in message.iter_attachments():
        name = part.get_filename()
        if name not in (XLSX, DOCX):
            continue
        if part.get("Content-Transfer-Encoding"):
            del part["Content-Transfer-Encoding"]
        part.set_payload((root / name).read_bytes())
        encoders.encode_base64(part)
    path.write_bytes(message.as_bytes())


def mutate_xlsx(root: Path, sheet: str, cell: str, value) -> None:
    path = root / XLSX
    workbook = load_workbook(path, data_only=False)
    workbook[sheet][cell] = value
    workbook.save(path)
    workbook.close()
    sync_email_attachments(root)


def replace_docx(root: Path, old: str, new: str, *, first_only: bool = False) -> None:
    path = root / DOCX
    document = Document(path)
    changed = False
    paragraphs = [*document.paragraphs, *[p for table in document.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]]
    for paragraph in paragraphs:
        if old in paragraph.text:
            paragraph.text = paragraph.text.replace(old, new, 1 if first_only else -1)
            changed = True
            if first_only:
                break
    if not changed:
        raise AssertionError(f"Word 中未找到待替换文本：{old}")
    document.save(path)
    sync_email_attachments(root)


def mutate_mail_body(root: Path, old: str, new: str) -> None:
    path = root / EML
    message = BytesParser(policy=policy.default.clone(linesep="\r\n")).parsebytes(path.read_bytes())
    changed = False
    for part in message.walk():
        if part.get_content_type() == "text/plain" and part.get_content_disposition() != "attachment":
            body = part.get_content()
            if old in body:
                part.set_content(body.replace(old, new, 1), subtype="plain", charset="utf-8")
                changed = True
                break
    if not changed:
        raise AssertionError(f"邮件正文中未找到待替换文本：{old}")
    path.write_bytes(message.as_bytes())


def one_source_value(root: Path) -> None:
    mutate_xlsx(root, "Reconciliation", "E6", 313.25)


def one_formula_flattened(root: Path) -> None:
    mutate_xlsx(root, "Reconciliation", "D2", 0)


def one_formula_wrong_reference(root: Path) -> None:
    mutate_xlsx(root, "Reconciliation", "L6", '=IF(J6="fail","fail",IF(ABS(G7)<0.005,"pass","fail"))')


def one_summary_formula_flattened(root: Path) -> None:
    mutate_xlsx(root, "Summary", "B3", 11)


def all_formulas_flattened(root: Path) -> None:
    path = root / XLSX
    formulas = load_workbook(path, data_only=False)
    values = load_workbook(path, data_only=True)
    for formula_sheet, value_sheet in zip(formulas.worksheets, values.worksheets, strict=True):
        for row in formula_sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    cell.value = value_sheet[cell.coordinate].value
    formulas.save(path)
    formulas.close()
    values.close()
    sync_email_attachments(root)


def missing_exception_row(root: Path) -> None:
    path = root / XLSX
    workbook = load_workbook(path, data_only=False)
    sheet = workbook["Exception Register"]
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 1).value == "WC-013":
            sheet.delete_rows(row, 1)
            break
    workbook.save(path)
    workbook.close()
    sync_email_attachments(root)


def exception_wrong_evidence(root: Path) -> None:
    path = root / XLSX
    workbook = load_workbook(path, data_only=False)
    sheet = workbook["Exception Register"]
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 1).value == "WC-005":
            sheet.cell(row, 3).value = "2024-01-15 | RUB 999.99 | Completed"
            break
    workbook.save(path)
    workbook.close()
    sync_email_attachments(root)


def exception_weak_action(root: Path) -> None:
    path = root / XLSX
    workbook = load_workbook(path, data_only=False)
    sheet = workbook["Exception Register"]
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 1).value == "WC-008":
            sheet.cell(row, 6).value = "Done"
            break
    workbook.save(path)
    workbook.close()
    sync_email_attachments(root)


def word_heading_demoted(root: Path) -> None:
    path = root / DOCX
    document = Document(path)
    target = next(p for p in document.paragraphs if p.text == "Evidence Boundary")
    target.style = document.styles["Normal"]
    document.save(path)
    sync_email_attachments(root)


def word_summary_conflict(root: Path) -> None:
    replace_docx(root, "11 orders matched", "12 orders matched", first_only=True)


def word_exception_missing(root: Path) -> None:
    replace_docx(root, "WC-013", "WC-999", first_only=True)


def email_summary_conflict(root: Path) -> None:
    mutate_mail_body(root, "11 matched", "12 matched")


def email_exception_omitted(root: Path) -> None:
    mutate_mail_body(root, "- WC-013: amount_mismatch", "")


def stale_email_attachment(root: Path) -> None:
    path = root / DOCX
    document = Document(path)
    document.add_paragraph("Internal review note.")
    document.save(path)


def ics_wrong_start(root: Path) -> None:
    path = root / ICS
    path.write_text(path.read_text(encoding="utf-8").replace("20240405T100000", "20240405T110000", 1), encoding="utf-8", newline="")


def ics_exception_omitted(root: Path) -> None:
    path = root / ICS
    path.write_text(path.read_text(encoding="utf-8").replace(", WC-013 amount_mismatch", "", 1), encoding="utf-8", newline="")


def ics_missing_uri(root: Path) -> None:
    path = root / ICS
    text = path.read_text(encoding="utf-8").replace(f"ATTACH;VALUE=URI:{DOCX}\n", "")
    path.write_text(text, encoding="utf-8", newline="")


def unsupported_fraud_claim(root: Path) -> None:
    path = root / DOCX
    document = Document(path)
    document.add_paragraph("WC-008 was confirmed as fraudulent.")
    document.save(path)
    sync_email_attachments(root)


def extra_debug_file(root: Path) -> None:
    (root / "debug.txt").write_text("temporary\n", encoding="utf-8")


PARTIAL = {
    "one_reconciliation_source_value_wrong": one_source_value,
    "one_reconciliation_formula_flattened": one_formula_flattened,
    "one_reconciliation_formula_wrong_reference": one_formula_wrong_reference,
    "one_summary_formula_flattened": one_summary_formula_flattened,
    "all_reconciliation_and_summary_formulas_flattened": all_formulas_flattened,
    "one_exception_row_missing": missing_exception_row,
    "one_exception_manual_evidence_wrong": exception_wrong_evidence,
    "one_exception_followup_too_weak": exception_weak_action,
    "word_one_heading_demoted": word_heading_demoted,
    "word_summary_conflict": word_summary_conflict,
    "word_one_exception_missing": word_exception_missing,
    "email_summary_conflict": email_summary_conflict,
    "email_one_exception_omitted": email_exception_omitted,
    "email_stale_docx_attachment": stale_email_attachment,
    "ics_wrong_start_time": ics_wrong_start,
    "ics_one_exception_omitted": ics_exception_omitted,
    "ics_one_relative_uri_missing": ics_missing_uri,
    "unsupported_fraud_assertion": unsupported_fraud_claim,
    "unexpected_debug_file": extra_debug_file,
}


def main() -> None:
    golden = JUDGE.evaluate(GOLDEN)
    assert golden["reward"] == 1.0, golden
    JUDGE._render_findings = lambda _artifact: [
        (True, "负例回归复用已验证 Golden 的 LibreOffice 可打开性"),
        (True, "负例回归复用已验证 Golden 的核心可见文本"),
        (True, "负例回归复用已验证 Golden 的 LibreOffice 可打开性"),
        (True, "负例回归复用已验证 Golden 的核心可见文本"),
    ]
    records = []
    with tempfile.TemporaryDirectory(prefix="financial-audit-v2-neg-") as raw:
        base = Path(raw)
        for name in REQUIRED:
            for mode in ("missing", "empty", "corrupt"):
                case = base / f"fatal-{mode}-{name}"
                copy_golden(case)
                target = case / name
                if mode == "missing":
                    target.unlink()
                elif mode == "empty":
                    target.write_bytes(b"")
                else:
                    target.write_bytes(b"\xff\xfe\x00")
                reward = JUDGE.evaluate(case)["reward"]
                assert reward == 0.0, (mode, name, reward)
                records.append({"case": f"{mode}:{name}", "class": "fatal", "reward": reward})
        case = base / "fatal-empty-output"
        case.mkdir()
        reward = JUDGE.evaluate(case)["reward"]
        assert reward == 0.0
        records.append({"case": "empty-output", "class": "fatal", "reward": reward})
        case = base / "fatal-input-only"
        shutil.copytree(ROOT / "data" / "input_files", case)
        reward = JUDGE.evaluate(case)["reward"]
        assert reward == 0.0
        records.append({"case": "input-only", "class": "fatal", "reward": reward})

        for name, mutate in PARTIAL.items():
            case = base / f"partial-{name}"
            copy_golden(case)
            mutate(case)
            result = JUDGE.evaluate(case)
            reward = float(result["reward"])
            assert 0.0 < reward < 1.0, (name, reward, result)
            records.append({"case": name, "class": "partial", "reward": reward})

    fatal = [record["reward"] for record in records if record["class"] == "fatal"]
    partial = [record["reward"] for record in records if record["class"] == "partial"]
    print(json.dumps({
        "fatal_count": len(fatal),
        "fatal_unique_rewards": sorted(set(fatal)),
        "partial_count": len(partial),
        "partial_min": min(partial),
        "partial_max": max(partial),
        "cases": records,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
