#!/usr/bin/env python3
"""按需构造并执行本题的任务专属负例，不在仓库中长期保存重复二进制。"""

from __future__ import annotations

import argparse
import shutil
import tempfile
from email import encoders, policy
from email.parser import BytesParser
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from judge import DOCX_NAME, EML_NAME, ICS_NAME, REQUIRED_FILES, XLSX_NAME, evaluate


CASES = (
    "missing-email",
    "formulas-flattened",
    "wrong-workbook-source-value",
    "wrong-word-summary",
    "stale-email-attachment",
    "wrong-calendar",
    "unsupported-fraud-claim",
    "email-summary-conflict",
    "corrupt-docx",
    "missing-exception-row",
    "extra-debug-file",
)


def copy_golden(golden: Path, case_dir: Path) -> None:
    case_dir.mkdir(parents=True, exist_ok=True)
    for name in REQUIRED_FILES:
        shutil.copy2(golden / name, case_dir / name)


def replace_docx_text(path: Path, old: str, new: str) -> None:
    document = Document(path)
    for paragraph in document.paragraphs:
        if old in paragraph.text:
            for run in paragraph.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if old in run.text:
                            run.text = run.text.replace(old, new)
    document.save(path)


def sync_email_attachments(case_dir: Path) -> None:
    path = case_dir / EML_NAME
    message = BytesParser(policy=policy.default.clone(linesep="\r\n")).parsebytes(path.read_bytes())
    for part in message.iter_attachments():
        name = part.get_filename()
        if name not in {XLSX_NAME, DOCX_NAME}:
            continue
        if part.get("Content-Transfer-Encoding"):
            del part["Content-Transfer-Encoding"]
        part.set_payload((case_dir / name).read_bytes())
        encoders.encode_base64(part)
    path.write_bytes(message.as_bytes())


def mutate(case: str, case_dir: Path) -> None:
    if case == "missing-email":
        (case_dir / EML_NAME).unlink()
    elif case == "formulas-flattened":
        path = case_dir / XLSX_NAME
        formulas = load_workbook(path, data_only=False)
        values = load_workbook(path, data_only=True)
        for formula_sheet, value_sheet in zip(formulas.worksheets, values.worksheets, strict=True):
            for row in formula_sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        cell.value = value_sheet[cell.coordinate].value
        formulas.save(path)
        sync_email_attachments(case_dir)
    elif case == "wrong-workbook-source-value":
        path = case_dir / XLSX_NAME
        workbook = load_workbook(path, data_only=False)
        sheet = workbook["Reconciliation"]
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row, 1).value == "WC-005":
                sheet.cell(row, 5).value = 999.99
                break
        workbook.save(path)
        sync_email_attachments(case_dir)
    elif case == "wrong-word-summary":
        replace_docx_text(case_dir / DOCX_NAME, "RUB 3,263.98", "RUB 3,200.00")
        sync_email_attachments(case_dir)
    elif case == "stale-email-attachment":
        document = Document(case_dir / DOCX_NAME)
        document.add_paragraph("Internal review copy.")
        document.save(case_dir / DOCX_NAME)
    elif case == "wrong-calendar":
        path = case_dir / ICS_NAME
        text = path.read_text(encoding="utf-8")
        text = text.replace("20240405T100000", "20240405T110000").replace("mailto:finance@company.com", "mailto:wrong@company.com")
        path.write_text(text, encoding="utf-8", newline="")
    elif case == "unsupported-fraud-claim":
        document = Document(case_dir / DOCX_NAME)
        document.add_paragraph("WC-008 was confirmed as fraudulent.")
        document.save(case_dir / DOCX_NAME)
        sync_email_attachments(case_dir)
    elif case == "email-summary-conflict":
        path = case_dir / EML_NAME
        message = BytesParser(policy=policy.default.clone(linesep="\r\n")).parsebytes(path.read_bytes())
        for part in message.walk():
            if part.get_content_type() == "text/plain" and part.get_content_disposition() != "attachment":
                body = part.get_content().replace("11 matched", "12 matched")
                part.set_content(body, subtype="plain", charset="utf-8")
                break
        path.write_bytes(message.as_bytes())
    elif case == "corrupt-docx":
        (case_dir / DOCX_NAME).write_bytes(b"not-a-docx")
    elif case == "missing-exception-row":
        path = case_dir / XLSX_NAME
        workbook = load_workbook(path, data_only=False)
        sheet = workbook["Exception Register"]
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row, 1).value == "WC-013":
                sheet.delete_rows(row, 1)
                break
        workbook.save(path)
        sync_email_attachments(case_dir)
    elif case == "extra-debug-file":
        (case_dir / "debug.txt").write_text("temporary", encoding="utf-8")
    else:
        raise ValueError(case)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--case", choices=CASES, action="append", required=True)
    args = parser.parse_args()
    failures = []
    for case in args.case:
        with tempfile.TemporaryDirectory(prefix=f"audit-negative-{case}-") as tmp:
            case_dir = Path(tmp)
            copy_golden(args.golden, case_dir)
            mutate(case, case_dir)
            result = evaluate(case_dir)
            print(f"{case}\treward={result['reward']:.6f}\tpassed={result['passed']}\thurdles={','.join(result['hurdle_failures'])}")
            if result["passed"] or float(result["reward"]) >= 0.8:
                failures.append({"case": case, "result": result})
    if failures:
        raise SystemExit(f"负例错误通过：{[item['case'] for item in failures]}")


if __name__ == "__main__":
    main()
