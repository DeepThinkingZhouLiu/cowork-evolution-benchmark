#!/usr/bin/env python3
"""根据本题冻结的持仓、MOEX 快照和投资政策生成三个相互一致的 Golden 交付物。"""

from __future__ import annotations

import argparse
import csv
import hashlib
import shutil
import subprocess
import tempfile
import zipfile
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, ROUND_HALF_UP
from email import policy
from email.message import EmailMessage
from email.parser import BytesParser
from email.utils import format_datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


XLSX_NAME = "Portfolio_Analysis.xlsx"
DOCX_NAME = "Investment_Memo.docx"
EML_NAME = "Monthly_Portfolio_Analysis.eml"
FIXED_DT = datetime(2026, 5, 27, 7, 30, 0, tzinfo=timezone.utc)
AS_OF = date(2026, 5, 26)
BASELINE_LIMIT = AS_OF - timedelta(days=30)
ALLOWED_SECTORS = {
    "Communication Services", "Consumer Defensive", "Financial Services", "Healthcare", "Energy"
}

NAVY = "17324D"
BLUE = "2A6F97"
TEAL = "168AAD"
GOLD = "D9A441"
WHITE = "FFFFFF"
PALE_BLUE = "EAF3F8"
PALE_GREEN = "E7F4EC"
PALE_RED = "FBE9E7"
PALE_GOLD = "FFF5DB"
LINE = "CCD8E2"
GRAY = "607080"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def money(value: Decimal) -> str:
    """货币展示使用常规财务的五入，底层计算仍保留全精度。"""
    return f"{value.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP):,.2f}"


def read_sources(input_dir: Path) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    portfolio_path = input_dir / "Portfolio.xlsx"
    snapshot_path = input_dir / "moex_market_snapshot.csv"
    policy_path = input_dir / "Investment_Policy.pdf"
    for path in (portfolio_path, snapshot_path, policy_path):
        if not path.is_file():
            raise FileNotFoundError(path)

    workbook = load_workbook(portfolio_path, data_only=False, read_only=True)
    if workbook.sheetnames != ["Holdings"]:
        raise ValueError(f"Portfolio.xlsx 工作表错误：{workbook.sheetnames}")
    sheet = workbook["Holdings"]
    headers = [str(cell.value or "").strip() for cell in sheet[1]]
    if headers != ["Symbol", "Company", "Quantity"]:
        raise ValueError(f"Portfolio.xlsx 表头错误：{headers}")
    holdings = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in row):
            continue
        holdings.append({"symbol": str(row[0]), "company": str(row[1]), "quantity": int(row[2])})
    workbook.close()
    if len(holdings) != 5 or len({item["symbol"] for item in holdings}) != 5:
        raise ValueError("Portfolio.xlsx 必须恰好包含 5 个不重复持仓")

    market_rows: list[dict[str, object]] = []
    with snapshot_path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        expected = ["Symbol", "Company", "Sector", "Currency", "Exchange", "Market_Cap", "Date", "Close"]
        if reader.fieldnames != expected:
            raise ValueError(f"moex_market_snapshot.csv 表头错误：{reader.fieldnames}")
        for row in reader:
            market_rows.append(
                {
                    "symbol": row["Symbol"], "company": row["Company"], "sector": row["Sector"],
                    "currency": row["Currency"], "exchange": row["Exchange"],
                    "market_cap": int(row["Market_Cap"]), "date": date.fromisoformat(row["Date"]),
                    "close": Decimal(row["Close"]),
                }
            )
    if len(market_rows) != 115:
        raise ValueError(f"市场快照必须恰好 115 行，实际 {len(market_rows)}")

    by_symbol: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in market_rows:
        by_symbol[str(row["symbol"])].append(row)
    if set(by_symbol) != {str(item["symbol"]) for item in holdings}:
        raise ValueError("市场快照与持仓代码集不一致")

    policy_text = subprocess.run(
        ["pdftotext", "-layout", str(policy_path), "-"],
        stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=60, check=True,
    ).stdout
    policy_text = " ".join(policy_text.split())
    for token in ("Communication Services", "Consumer Defensive", "Financial Services", "Healthcare", "Energy", "RUB", "MCX", "100 млрд"):
        if token not in policy_text:
            raise ValueError(f"Investment_Policy.pdf 缺少关键政策：{token}")

    analyzed: list[dict[str, object]] = []
    for holding in holdings:
        symbol = str(holding["symbol"])
        history = sorted(by_symbol[symbol], key=lambda item: item["date"])
        latest = max(history, key=lambda item: item["date"])
        baseline = max((item for item in history if item["date"] <= BASELINE_LIMIT), key=lambda item: item["date"])
        if latest["date"] != AS_OF or baseline["date"] != date(2026, 4, 24):
            raise ValueError(f"{symbol} 行情边界不符合冻结协议")
        issuer_facts = {
            (str(item["company"]), str(item["sector"]), str(item["currency"]), str(item["exchange"]), int(item["market_cap"]))
            for item in history
        }
        if len(issuer_facts) != 1:
            raise ValueError(f"{symbol} 发行人字段在快照中不一致")
        company, sector, currency, exchange, market_cap = issuer_facts.pop()
        if company != holding["company"]:
            raise ValueError(f"{symbol} 公司名与 Portfolio.xlsx 冲突")
        quantity = int(holding["quantity"])
        current_value = Decimal(quantity) * Decimal(latest["close"])
        return_pct = Decimal(latest["close"]) / Decimal(baseline["close"]) - Decimal("1")
        analyzed.append(
            {
                **holding, "sector": sector, "baseline_date": baseline["date"], "baseline_close": baseline["close"],
                "current_date": latest["date"], "current_price": latest["close"], "current_value": current_value,
                "return_pct": return_pct, "market_cap": market_cap, "currency": currency, "exchange": exchange,
            }
        )
    total = sum((Decimal(item["current_value"]) for item in analyzed), Decimal("0"))
    for item in analyzed:
        item["weight_pct"] = Decimal(item["current_value"]) / total
        item["market_cap_category"] = "Large Cap" if int(item["market_cap"]) > 100_000_000_000 else "Below Large Cap"
        item["policy_status"] = (
            "Compliant" if item["sector"] in ALLOWED_SECTORS and item["currency"] == "RUB"
            and item["exchange"] == "MCX" and item["market_cap_category"] == "Large Cap" else "Review"
        )
    return analyzed, market_rows


def style_sheet(sheet, widths: dict[str, float], orientation: str = "landscape", paper_size: int | None = None) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 34
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", size=8, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color=LINE)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=8, color=NAVY)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.page_setup.orientation = orientation
    sheet.page_setup.paperSize = paper_size or sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:1"
    sheet.page_margins.left = 0.2
    sheet.page_margins.right = 0.2
    sheet.page_margins.top = 0.35
    sheet.page_margins.bottom = 0.35


def add_native_table(sheet, name: str) -> None:
    table = Table(displayName=name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    sheet.add_table(table)


def build_workbook(analyzed: list[dict[str, object]], market_rows: list[dict[str, object]], output_path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.properties.creator = "EvoCowork benchmark adapter"
    workbook.properties.title = "MOEX Portfolio Analysis"
    workbook.properties.created = FIXED_DT.replace(tzinfo=None)
    workbook.properties.modified = FIXED_DT.replace(tzinfo=None)

    holdings = workbook.create_sheet("Holdings")
    headers = [
        "Symbol", "Company", "Sector", "Quantity", "Baseline_Date", "Baseline_Close", "Current_Date",
        "Current_Price", "Current_Value", "Return_30d_Pct", "Weight_Pct", "Market_Cap_RUB",
        "Market_Cap_Category", "Currency", "Exchange", "Policy_Status",
    ]
    holdings.append(headers)
    allowed_formula = ",".join(f'"{value}"' for value in sorted(ALLOWED_SECTORS))
    for excel_row, item in enumerate(analyzed, start=2):
        holdings.append(
            [
                item["symbol"], item["company"], item["sector"], item["quantity"], item["baseline_date"],
                float(item["baseline_close"]), item["current_date"], float(item["current_price"]),
                f"=D{excel_row}*H{excel_row}", f"=H{excel_row}/F{excel_row}-1",
                f"=I{excel_row}/SUM($I$2:$I$6)", item["market_cap"],
                f'=IF(L{excel_row}>100000000000,"Large Cap","Below Large Cap")',
                item["currency"], item["exchange"],
                f'=IF(AND(OR(C{excel_row}={allowed_formula.replace(",", f",C{excel_row}=")}),N{excel_row}="RUB",O{excel_row}="MCX",M{excel_row}="Large Cap"),"Compliant","Review")',
            ]
        )
    style_sheet(
        holdings,
        {"A": 13, "B": 19, "C": 22, "D": 9, "E": 13, "F": 14, "G": 13, "H": 14,
         "I": 15, "J": 14, "K": 12, "L": 18, "M": 17, "N": 10, "O": 10, "P": 14},
        paper_size=holdings.PAPERSIZE_A3,
    )
    for row in range(2, 7):
        holdings.row_dimensions[row].height = 34
        for column in (5, 7):
            holdings.cell(row, column).number_format = "yyyy-mm-dd"
        for column in (6, 8, 9):
            holdings.cell(row, column).number_format = '#,##0.00 "RUB"'
        for column in (10, 11):
            holdings.cell(row, column).number_format = "0.00%"
        holdings.cell(row, 12).number_format = '#,##0 "RUB"'
    holdings.conditional_formatting.add("J2:J6", FormulaRule(formula=["J2<0"], fill=PatternFill("solid", fgColor=PALE_RED)))
    holdings.conditional_formatting.add("P2:P6", FormulaRule(formula=['P2="Compliant"'], fill=PatternFill("solid", fgColor=PALE_GREEN)))
    add_native_table(holdings, "PortfolioHoldingsTable")
    holdings.print_area = "A1:P6"
    holdings.page_setup.fitToHeight = 1

    summary = workbook.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    metrics = [
        ("Total_Portfolio_Value", "=SUM(Holdings!I2:I6)"),
        ("Best_Performer", "=INDEX(Holdings!A2:A6,MATCH(MAX(Holdings!J2:J6),Holdings!J2:J6,0))"),
        ("Worst_Performer", "=INDEX(Holdings!A2:A6,MATCH(MIN(Holdings!J2:J6),Holdings!J2:J6,0))"),
        ("Financial Services", '=SUMIF(Holdings!C2:C6,A5,Holdings!I2:I6)'),
        ("Energy", '=SUMIF(Holdings!C2:C6,A6,Holdings!I2:I6)'),
        ("Consumer Defensive", '=SUMIF(Holdings!C2:C6,A7,Holdings!I2:I6)'),
        ("Communication Services", '=SUMIF(Holdings!C2:C6,A8,Holdings!I2:I6)'),
        ("Largest_Sector", "=INDEX(A5:A8,MATCH(MAX(B5:B8),B5:B8,0))"),
        ("Largest_Sector_Weight", "=MAX(B5:B8)/B2"),
        ("Policy_Compliant_Positions", '=COUNTIF(Holdings!P2:P6,"Compliant")'),
    ]
    for label, formula in metrics:
        summary.append([label, formula])
    style_sheet(summary, {"A": 34, "B": 24}, orientation="landscape")
    summary.page_setup.fitToHeight = 1
    for row in range(2, summary.max_row + 1):
        summary.cell(row, 1).fill = PatternFill("solid", fgColor=PALE_BLUE)
        summary.cell(row, 1).font = Font(name="Aptos", size=10, bold=True, color=NAVY)
        summary.cell(row, 2).font = Font(name="Aptos Display", size=12, bold=True, color=TEAL)
    for row in (2, 5, 6, 7, 8):
        summary.cell(row, 2).number_format = '#,##0.00 "RUB"'
    summary.cell(10, 2).number_format = "0.00%"
    summary.conditional_formatting.add("B10", FormulaRule(formula=["B10>0.4"], fill=PatternFill("solid", fgColor=PALE_GOLD)))
    add_native_table(summary, "PortfolioSummaryTable")
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "Sector exposure (RUB)"
    chart.y_axis.title = "Sector"
    chart.x_axis.title = "RUB"
    chart.x_axis.numFmt = '0,"K"'
    chart.x_axis.majorUnit = 100000
    chart.add_data(Reference(summary, min_col=2, min_row=5, max_row=8), titles_from_data=False)
    chart.set_categories(Reference(summary, min_col=1, min_row=5, max_row=8))
    chart.height = 6.2
    chart.width = 12.0
    chart.legend = None
    summary.add_chart(chart, "D2")
    summary.print_area = "A1:L11"

    source = workbook.create_sheet("Source_Prices")
    source.append(["Symbol", "Company", "Sector", "Currency", "Exchange", "Market_Cap", "Date", "Close"])
    for row in market_rows:
        source.append(
            [row["symbol"], row["company"], row["sector"], row["currency"], row["exchange"], row["market_cap"], row["date"], float(row["close"])]
        )
    style_sheet(source, {"A": 13, "B": 20, "C": 23, "D": 10, "E": 10, "F": 18, "G": 13, "H": 14})
    for row in range(2, source.max_row + 1):
        source.cell(row, 6).number_format = '#,##0 "RUB"'
        source.cell(row, 7).number_format = "yyyy-mm-dd"
        source.cell(row, 8).number_format = '#,##0.0000 "RUB"'
    add_native_table(source, "SourcePricesTable")
    source.print_area = f"A1:H{source.max_row}"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    recalculate_workbook(output_path)
    restore_frozen_header_rows(output_path, sheet_count=3)
    normalize_ooxml(output_path)


def recalculate_workbook(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="portfolio-xlsx-recalc-") as tmp:
        root = Path(tmp)
        input_root = root / "input"
        output_root = root / "output"
        profile = root / "profile"
        input_root.mkdir()
        output_root.mkdir()
        profile.mkdir()
        source = input_root / path.name
        shutil.copy2(path, source)
        result = subprocess.run(
            ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "xlsx", "--outdir", str(output_root), str(source)],
            stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300,
        )
        converted = output_root / path.name
        if result.returncode != 0 or not converted.is_file():
            raise RuntimeError(f"LibreOffice 公式重算失败：{result.stdout[-1000:]}")
        shutil.copy2(converted, path)


def restore_frozen_header_rows(path: Path, sheet_count: int) -> None:
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with tempfile.TemporaryDirectory(prefix="portfolio-xlsx-freeze-") as tmp:
        target = Path(tmp) / path.name
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w") as destination:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename in {f"xl/worksheets/sheet{index}.xml" for index in range(1, sheet_count + 1)}:
                    root = etree.fromstring(data)
                    sheet_view = root.find(f"{{{namespace}}}sheetViews/{{{namespace}}}sheetView")
                    if sheet_view is not None:
                        for existing in sheet_view.findall(f"{{{namespace}}}pane"):
                            sheet_view.remove(existing)
                        pane = etree.Element(f"{{{namespace}}}pane")
                        pane.set("ySplit", "1")
                        pane.set("topLeftCell", "A2")
                        pane.set("activePane", "bottomLeft")
                        pane.set("state", "frozen")
                        sheet_view.insert(0, pane)
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                destination.writestr(info, data)
        shutil.copy2(target, path)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE, size: str = "6") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def add_word_table(document: Document, headers: list[str], data: list[list[str]], widths: list[float]):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, NAVY)
        set_cell_border(cell, NAVY)
        cell.width = Cm(widths[index])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(7.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(data):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.text = str(value)
            set_cell_shading(cell, WHITE if row_index % 2 == 0 else "F4F7FA")
            set_cell_border(cell)
            cell.width = Cm(widths[index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(7.5)
                    run.font.color.rgb = RGBColor.from_string(NAVY)
    return table


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("MOEX Portfolio  |  ")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)


def build_memo(analyzed: list[dict[str, object]], output_path: Path) -> None:
    total = sum((Decimal(item["current_value"]) for item in analyzed), Decimal("0"))
    best = max(analyzed, key=lambda item: Decimal(item["return_pct"]))
    worst = min(analyzed, key=lambda item: Decimal(item["return_pct"]))
    sectors: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    for item in analyzed:
        sectors[str(item["sector"])] += Decimal(item["current_value"])
    largest_sector = max(sectors, key=sectors.get)
    largest_value = sectors[largest_sector]
    largest_weight = largest_value / total

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    document.core_properties.creator = "EvoCowork benchmark adapter"
    document.core_properties.title = "Portfolio Analysis Report"
    document.core_properties.created = FIXED_DT.replace(tzinfo=None)
    document.core_properties.modified = FIXED_DT.replace(tzinfo=None)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].font.color.rgb = RGBColor.from_string(NAVY)
    for name in ("Title", "Heading 1", "Heading 2"):
        styles[name].font.name = "Aptos Display"
        styles[name].font.color.rgb = RGBColor.from_string(NAVY)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Portfolio Analysis Report")
    run.font.size = Pt(28)
    run.font.bold = True
    subtitle = document.add_paragraph()
    subtitle_run = subtitle.add_run("MOEX holdings  |  Valuation date: 26 May 2026  |  Currency: RUB")
    subtitle_run.font.name = "Aptos"
    subtitle_run.font.size = Pt(10)
    subtitle_run.font.color.rgb = RGBColor.from_string(TEAL)
    document.add_paragraph(
        "Prepared for the Investment Committee from the frozen portfolio, market snapshot and Fund Investment Policy v2026.1."
    )

    add_word_table(
        document,
        ["Portfolio value", "Best 30d", "Worst 30d", "Largest sector"],
        [[f"{money(total)} RUB", f"{best['symbol']}  {Decimal(best['return_pct']):.2%}", f"{worst['symbol']}  {Decimal(worst['return_pct']):.2%}", f"{largest_sector}  {largest_weight:.2%}"]],
        [4.5, 4.5, 4.5, 4.8],
    )

    document.add_heading("Holdings Overview", level=1)
    document.add_paragraph(
        "All values use the latest close on 26 May 2026. The 30-day comparison uses the latest close on or before 26 April 2026, which is 24 April 2026 for every holding."
    )
    holdings_rows = []
    for item in analyzed:
        holdings_rows.append(
            [
                str(item["symbol"]), str(item["company"]), str(item["sector"]), f"{int(item['quantity']):,}",
                f"{Decimal(item['current_price']):,.4f}", money(Decimal(item["current_value"])),
                f"{Decimal(item['return_pct']):.2%}", f"{Decimal(item['weight_pct']):.2%}",
            ]
        )
    add_word_table(
        document,
        ["Symbol", "Company", "Sector", "Qty", "Price", "Value (RUB)", "30d", "Weight"],
        holdings_rows,
        [2.0, 2.7, 3.4, 1.2, 2.0, 2.4, 1.5, 1.5],
    )

    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_heading("Performance Analysis", level=1)
    document.add_paragraph(
        f"Total portfolio value is {money(total)} RUB. {best['symbol']} is the best 30-day performer at {Decimal(best['return_pct']):.2%}; "
        f"{worst['symbol']} is the worst at {Decimal(worst['return_pct']):.2%}. These figures describe price changes in the frozen window and do not establish their causes."
    )
    sector_rows = []
    for sector_name, value in sorted(sectors.items(), key=lambda item: item[1], reverse=True):
        sector_rows.append([sector_name, f"{money(value)} RUB", f"{value / total:.2%}"])
    add_word_table(document, ["Sector", "Market value", "Portfolio weight"], sector_rows, [7.4, 5.5, 4.2])
    document.add_paragraph(
        f"{largest_sector} is the largest sector exposure at {money(largest_value)} RUB ({largest_weight:.2%}). "
        "GAZP.ME and LKOH.ME are combined in this sector total."
    )

    document.add_heading("Policy Review", level=1)
    policy_items = [
        "All five issuers are in sectors permitted by the policy.",
        "All quoted prices and portfolio values are in RUB and the exchange is MCX.",
        "All five issuers have market capitalisation above 100 billion RUB and are classified as Large Cap.",
        "The supplied policy does not define a maximum position or sector-weight limit; this report therefore does not invent one.",
    ]
    for item in policy_items:
        document.add_paragraph(item, style="List Bullet")
    document.add_paragraph("Policy-compliant positions: 5 of 5.")

    document.add_heading("Recommendations", level=1)
    recommendations = [
        f"Monitor SBER.ME before the next committee meeting because it has the weakest observed 30-day return ({Decimal(worst['return_pct']):.2%}); review new evidence rather than inferring a cause from price movement alone.",
        f"Review the {largest_sector} exposure ({largest_weight:.2%}) and decide whether the policy should introduce an explicit sector-concentration threshold; no such threshold exists in the supplied policy.",
        f"Continue tracking MTSS.ME after its positive 30-day return ({Decimal(best['return_pct']):.2%}), without treating one window as a forecast or an automatic buy signal.",
    ]
    for item in recommendations:
        document.add_paragraph(item, style="List Number")

    document.add_heading("Method and Evidence Boundary", level=1)
    document.add_paragraph(
        "Source files: Portfolio.xlsx, moex_market_snapshot.csv and Investment_Policy.pdf. Current values equal quantity multiplied by the latest close. Returns compare the latest close with the latest close on or before the 30-calendar-day cutoff. No external market data or unsupported investment recommendation was used."
    )
    # 新分节默认链接前一节页脚；只写一次，避免重复页码字段。
    add_page_number(document.sections[0].footer.paragraphs[0])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    normalize_ooxml(output_path)


def build_email(analyzed: list[dict[str, object]], xlsx_path: Path, docx_path: Path, output_path: Path) -> None:
    total = sum((Decimal(item["current_value"]) for item in analyzed), Decimal("0"))
    best = max(analyzed, key=lambda item: Decimal(item["return_pct"]))
    worst = min(analyzed, key=lambda item: Decimal(item["return_pct"]))
    message = EmailMessage(policy=policy.default.clone(linesep="\r\n"))
    message["Date"] = format_datetime(FIXED_DT)
    message["From"] = "analyst@fund.example.com"
    message["To"] = "investment-committee@fund.example.com"
    message["Subject"] = "Monthly Portfolio Analysis Report"
    message["Message-ID"] = "<monthly-portfolio-analysis-20260526@fund.example.com>"
    message.set_content(
        "Investment Committee,\n\n"
        f"The portfolio value as of 2026-05-26 is {money(total)} RUB.\n"
        f"Best 30-day performer: {best['symbol']} ({Decimal(best['return_pct']):.2%}).\n"
        f"Worst 30-day performer: {worst['symbol']} ({Decimal(worst['return_pct']):.2%}).\n\n"
        "The attached workbook contains the formula-based analysis and frozen source prices. "
        "The attached memo contains the committee summary, policy review and evidence-bounded recommendations.\n\n"
        "Regards,\nPortfolio Analytics",
        subtype="plain", charset="utf-8",
    )
    message.add_attachment(
        xlsx_path.read_bytes(), maintype="application", subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=XLSX_NAME,
    )
    message.add_attachment(
        docx_path.read_bytes(), maintype="application", subtype="vnd.openxmlformats-officedocument.wordprocessingml.document", filename=DOCX_NAME,
    )
    message.set_boundary("===============EvoCoworkPortfolio20260526==")
    output_path.write_bytes(message.as_bytes())
    parsed = BytesParser(policy=policy.default).parsebytes(output_path.read_bytes())
    attachments = list(parsed.iter_attachments())
    if [part.get_filename() for part in attachments] != [XLSX_NAME, DOCX_NAME]:
        raise RuntimeError("生成的 EML 附件不完整")
    if sha256(xlsx_path) != hashlib.sha256(attachments[0].get_payload(decode=True)).hexdigest():
        raise RuntimeError("EML 中的 Excel 附件与实际交付不同")
    if sha256(docx_path) != hashlib.sha256(attachments[1].get_payload(decode=True)).hexdigest():
        raise RuntimeError("EML 中的 Word 附件与实际交付不同")


def normalize_ooxml(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="portfolio-ooxml-normalize-") as tmp:
        target = Path(tmp) / path.name
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w") as destination:
            for info in sorted(source.infolist(), key=lambda item: item.filename):
                data = source.read(info.filename)
                if info.filename == "docProps/core.xml":
                    root = etree.fromstring(data)
                    dcterms = "http://purl.org/dc/terms/"
                    for field in ("created", "modified"):
                        node = root.find(f"{{{dcterms}}}{field}")
                        if node is not None:
                            node.text = "2026-05-27T07:30:00Z"
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                elif info.filename.startswith("xl/charts/") and info.filename.endswith(".xml"):
                    # LibreOffice 每次重算会随机生成图表轴 ID；把同一图中的分类轴和值轴
                    # 映射到固定 ID，确保 fresh replay 与邮件附件能够逐字节复现。
                    root = etree.fromstring(data)
                    chart_ns = "http://schemas.openxmlformats.org/drawingml/2006/chart"
                    category_axis = root.find(f".//{{{chart_ns}}}catAx/{{{chart_ns}}}axId")
                    value_axis = root.find(f".//{{{chart_ns}}}valAx/{{{chart_ns}}}axId")
                    if category_axis is not None and value_axis is not None:
                        mapping = {
                            category_axis.get("val", ""): "11084706",
                            value_axis.get("val", ""): "27695691",
                        }
                        for node in root.iter():
                            if etree.QName(node).localname in {"axId", "crossAx"} and node.get("val") in mapping:
                                node.set("val", mapping[node.get("val")])
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                copied = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
                copied.compress_type = zipfile.ZIP_DEFLATED
                copied.external_attr = info.external_attr
                copied.create_system = info.create_system
                copied.flag_bits = info.flag_bits
                destination.writestr(copied, data)
        shutil.copy2(target, path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for child in args.output_dir.iterdir():
        if child.is_file():
            child.unlink()
        elif child.is_dir():
            shutil.rmtree(child)
    analyzed, market_rows = read_sources(args.input_dir)
    xlsx_path = args.output_dir / XLSX_NAME
    docx_path = args.output_dir / DOCX_NAME
    build_workbook(analyzed, market_rows, xlsx_path)
    build_memo(analyzed, docx_path)
    build_email(analyzed, xlsx_path, docx_path, args.output_dir / EML_NAME)


if __name__ == "__main__":
    main()
