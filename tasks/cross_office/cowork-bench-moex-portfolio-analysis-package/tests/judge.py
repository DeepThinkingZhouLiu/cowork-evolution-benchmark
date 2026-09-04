#!/usr/bin/env python3
"""MOEX 月度投资组合分析包的 v2 连续评分入口；v1 原样保存在 judge_v1.py。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import re
import shutil
import sys
from decimal import Decimal
from email import policy
from email.parser import BytesParser
from pathlib import Path
from typing import Iterable

from docx import Document
from openpyxl import load_workbook


TASK_ID = "cowork-evo/cowork-bench-moex-portfolio-analysis-package"
RUBRIC_VERSION = "cowork-bench-moex-portfolio-analysis-v0.2"
JUDGE_VERSION = "cowork-bench-moex-portfolio-analysis-criterion-partial-v2.0"
XLSX = "Portfolio_Analysis.xlsx"
DOCX = "Investment_Memo.docx"
EML = "Monthly_Portfolio_Analysis.eml"
REQUIRED = (XLSX, DOCX, EML)
POSITIVE_WEIGHTS = {
    "R001": 5, "R002": 8, "R003": 8, "R004": 12, "R005": 7, "R006": 10,
    "R007": 10, "R008": 9, "R009": 7, "R010": 5, "R011": 6, "R012": 7, "R013": 6,
}
PENALTY_WEIGHTS = {"R014": -30, "R015": -25, "R016": -30, "R017": -25, "R018": -25}
SUMMARY_TOKENS = (
    ("SUM", "HOLDINGS", "I2:I6"),
    ("INDEX", "MATCH", "MAX", "HOLDINGS", "A2:A6", "J2:J6"),
    ("INDEX", "MATCH", "MIN", "HOLDINGS", "A2:A6", "J2:J6"),
    ("SUMIF", "HOLDINGS", "C2:C6", "A5", "I2:I6"),
    ("SUMIF", "HOLDINGS", "C2:C6", "A6", "I2:I6"),
    ("SUMIF", "HOLDINGS", "C2:C6", "A7", "I2:I6"),
    ("SUMIF", "HOLDINGS", "C2:C6", "A8", "I2:I6"),
    ("INDEX", "MATCH", "MAX", "A5:A8", "B5:B8"),
    ("MAX", "B5:B8", "B2"),
    ("COUNTIF", "HOLDINGS", "P2:P6", "COMPLIANT"),
)


def _load_v1():
    path = Path(__file__).with_name("judge_v1.py")
    if str(path.parent) not in sys.path:
        sys.path.insert(0, str(path.parent))
    spec = importlib.util.spec_from_file_location("moex_portfolio_judge_v1", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v1 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


V1 = _load_v1()


def _sha(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _norm(value: object) -> str:
    return V1.normalize(value)


def _lower(value: object) -> str:
    return _norm(value).lower()


def _approx(actual: object, expected: object, tolerance: float = 1e-7) -> bool:
    try:
        return abs(float(actual) - float(expected)) <= tolerance
    except (TypeError, ValueError):
        return False


def _same(actual: object, expected: object) -> bool:
    if isinstance(expected, (Decimal, int, float)) and not isinstance(expected, bool):
        return _approx(actual, expected)
    return _norm(actual) == _norm(expected)


def _formula(cell, *tokens: str) -> bool:
    text = _norm(cell.value).upper().replace("$", "")
    return cell.data_type == "f" and all(token.upper().replace("$", "") in text for token in tokens)


def _ratio(findings: Iterable[tuple[bool, str]]) -> tuple[float, list[str], int, int]:
    rows = list(findings)
    earned = sum(bool(ok) for ok, _ in rows)
    return earned / max(1, len(rows)), [label for ok, label in rows if not ok], earned, len(rows)


def _criterion(key: str, score: float, evidence: object, earned: float = 0, total: float = 1):
    return {
        "criterion_id": key, "score": round(max(0.0, min(1.0, float(score))), 6),
        "evidence": evidence, "method": "deterministic",
        "raw": {"earned_units": round(earned, 6), "total_units": round(total, 6)},
    }


def _message(path: Path):
    return BytesParser(policy=policy.default).parsebytes(path.read_bytes())


def _mail_body(message) -> str:
    try:
        body = message.get_body(preferencelist=("plain",)) if message.is_multipart() else message
        return body.get_content() if body else ""
    except Exception:
        return ""


def _core(artifact: Path):
    errors: list[str] = []
    hashes: list[str] = []
    truth, source_rows, source_errors, summary = V1.load_truth()
    if source_errors or not truth or len(source_rows) != 115:
        errors.append(f"冻结输入校验失败：{source_errors}")
    if not artifact.is_dir():
        return False, ["输出目录不存在", *errors], hashes, truth, source_rows, summary
    source_hashes = set(V1.SOURCE_HASHES.values())
    for name in REQUIRED:
        path = artifact / name
        if not path.is_file() or path.stat().st_size == 0:
            errors.append(f"核心交付缺失或为空：{name}")
            continue
        digest = _sha(path)
        hashes.append(digest)
        if digest in source_hashes:
            errors.append(f"核心交付只是原样输入：{name}")
            continue
        try:
            if name == XLSX:
                workbook = load_workbook(path, read_only=True, data_only=False)
                if not workbook.sheetnames or all(sheet.max_row == 0 for sheet in workbook.worksheets):
                    raise ValueError("工作簿没有内容")
                workbook.close()
            elif name == DOCX:
                document = Document(path)
                if not document.paragraphs and not document.tables:
                    raise ValueError("Word 没有内容")
            else:
                message = _message(path)
                if not list(message.items()) or not message.get_payload():
                    raise ValueError("EML 没有可解析头部或正文")
        except Exception as exc:  # noqa: BLE001
            errors.append(f"核心交付损坏或不可解析：{name}: {type(exc).__name__}: {exc}")
    return not errors, errors, hashes, truth, source_rows, summary


def _fatal(errors: list[str], hashes: list[str]):
    return {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes or ["0" * 64],
        "criterion_results": [_criterion(key, 0, "；".join(errors), 0, 0) for key in (*POSITIVE_WEIGHTS, *PENALTY_WEIGHTS)],
        "reward": 0.0, "passed": False, "hurdle_failures": ["R001"], "penalties": [],
    }


def _calculated_findings(path: Path, truth, summary):
    findings: dict[str, list[tuple[bool, str]]] = {"calculations": [], "policy": [], "summary": []}
    recalculated, log = V1.recalculate_candidate(path)
    if recalculated is None:
        failure = (False, f"LibreOffice 公式重算失败：{log}")
        return {key: [failure] for key in findings}
    try:
        workbook = load_workbook(recalculated, data_only=True)
        holdings = workbook["Holdings"]
        for row, expected in enumerate(truth, start=2):
            for column, wanted in zip((9, 10, 11), (expected["current_value"], expected["return_pct"], expected["weight_pct"]), strict=True):
                findings["calculations"].append((_approx(holdings.cell(row, column).value, wanted), f"Holdings!{holdings.cell(row, column).coordinate} 重算值错误"))
            findings["policy"].append((_norm(holdings.cell(row, 13).value) == expected["market_cap_category"], f"Holdings!M{row} 重算分类错误"))
            findings["policy"].append((_norm(holdings.cell(row, 16).value) == expected["policy_status"], f"Holdings!P{row} 重算政策状态错误"))
        sector_values = summary["sectors"]
        expected_summary = [
            summary["total"], summary["best"]["symbol"], summary["worst"]["symbol"],
            sector_values["Financial Services"], sector_values["Energy"], sector_values["Consumer Defensive"],
            sector_values["Communication Services"], summary["largest_sector"], summary["largest_sector_weight"], summary["compliant"],
        ]
        sheet = workbook["Summary"]
        for row, wanted in enumerate(expected_summary, start=2):
            findings["summary"].append((_same(sheet.cell(row, 2).value, wanted), f"Summary!B{row} 重算值错误"))
        workbook.close()
    except Exception as exc:  # noqa: BLE001
        failure = (False, f"重算工作簿解析失败：{type(exc).__name__}: {exc}")
        findings = {key: [failure] for key in findings}
    finally:
        shutil.rmtree(recalculated.parents[1], ignore_errors=True)
    return findings


def _workbook(path: Path, truth, source_rows, summary):
    workbook = load_workbook(path, data_only=False)
    identity: list[tuple[bool, str]] = []
    boundary: list[tuple[bool, str]] = []
    calculations: list[tuple[bool, str]] = []
    policy: list[tuple[bool, str]] = []
    summary_checks: list[tuple[bool, str]] = []
    source_checks: list[tuple[bool, str]] = []
    native: list[tuple[bool, str]] = []
    formula_count = 0
    formula_total = 35
    delivered_symbols: list[str] = []

    identity.append((workbook.sheetnames == ["Holdings", "Summary", "Source_Prices"], f"工作表顺序/集合错误：{workbook.sheetnames}"))
    if "Holdings" in workbook.sheetnames:
        sheet = workbook["Holdings"]
        for column, header in enumerate(V1.HOLDING_HEADERS, start=1):
            identity.append((_norm(sheet.cell(1, column).value) == header, f"Holdings 表头 {column} 错误"))
        identity.append((sheet.max_row == 6 and sheet.max_column == 16, f"Holdings 形状错误：{sheet.max_row}x{sheet.max_column}"))
        for row, expected in enumerate(truth, start=2):
            delivered_symbols.append(_norm(sheet.cell(row, 1).value))
            static = {
                1: expected["symbol"], 2: expected["company"], 3: expected["sector"], 4: expected["quantity"],
                12: expected["market_cap"], 14: expected["currency"], 15: expected["exchange"],
            }
            for column, wanted in static.items():
                identity.append((_same(sheet.cell(row, column).value, wanted), f"Holdings!{sheet.cell(row, column).coordinate} 发行人/持仓值错误"))
            for column, wanted, kind in (
                (5, expected["baseline_date"].isoformat(), "date"), (6, expected["baseline_close"], "number"),
                (7, expected["current_date"].isoformat(), "date"), (8, expected["current_price"], "number"),
            ):
                actual = sheet.cell(row, column).value
                ok = V1.iso_date(actual) == wanted if kind == "date" else _approx(actual, wanted)
                boundary.append((ok, f"Holdings!{sheet.cell(row, column).coordinate} 行情边界错误"))
            formula_rules = {
                9: (f"D{row}*H{row}",),
                10: (f"H{row}/F{row}-1",),
                11: (f"I{row}/SUM(I2:I6)",),
                13: (f"L{row}>100000000000", "LARGE CAP", "BELOW LARGE CAP"),
                16: ("IF", "AND", f"C{row}", f"N{row}", f"O{row}", f"M{row}", "RUB", "MCX", "LARGE CAP", "COMPLIANT", "REVIEW"),
            }
            for column in (9, 10, 11):
                cell = sheet.cell(row, column)
                calculations.append((_formula(cell, *formula_rules[column]), f"Holdings!{cell.coordinate} 计算公式错误"))
                formula_count += int(cell.data_type == "f")
            for column in (13, 16):
                cell = sheet.cell(row, column)
                policy.append((_formula(cell, *formula_rules[column]), f"Holdings!{cell.coordinate} 政策公式错误"))
                formula_count += int(cell.data_type == "f")
    else:
        identity.append((False, "缺少 Holdings"))
        boundary.append((False, "缺少 Holdings"))
        calculations.append((False, "缺少 Holdings"))
        policy.append((False, "缺少 Holdings"))

    recalculated = _calculated_findings(path, truth, summary)
    calculations.extend(recalculated["calculations"])
    policy.extend(recalculated["policy"])

    if "Summary" in workbook.sheetnames:
        sheet = workbook["Summary"]
        summary_checks.extend([(_norm(sheet["A1"].value) == "Metric", "Summary!A1 错误"), (_norm(sheet["B1"].value) == "Value", "Summary!B1 错误")])
        for row, label in enumerate(V1.SUMMARY_LABELS, start=2):
            summary_checks.append((_norm(sheet.cell(row, 1).value) == label, f"Summary!A{row} 标签错误"))
            cell = sheet.cell(row, 2)
            summary_checks.append((_formula(cell, *SUMMARY_TOKENS[row - 2]), f"Summary!B{row} 公式错误"))
            formula_count += int(cell.data_type == "f")
        summary_checks.extend(recalculated["summary"])
    else:
        summary_checks.append((False, "缺少 Summary"))

    if "Source_Prices" in workbook.sheetnames:
        sheet = workbook["Source_Prices"]
        for column, header in enumerate(V1.SOURCE_HEADERS, start=1):
            source_checks.append((_norm(sheet.cell(1, column).value) == header, f"Source_Prices 表头 {column} 错误"))
        source_checks.append((sheet.max_row == 116 and sheet.max_column == 8, f"Source_Prices 形状错误：{sheet.max_row}x{sheet.max_column}"))
        for row, expected in enumerate(source_rows, start=2):
            wanted = [
                expected["symbol"], expected["company"], expected["sector"], expected["currency"],
                expected["exchange"], expected["market_cap"], expected["date"].isoformat(), expected["close"],
            ]
            for column, target in enumerate(wanted, start=1):
                actual = sheet.cell(row, column).value
                if column == 7:
                    ok = V1.iso_date(actual) == target
                else:
                    ok = _same(actual, target)
                source_checks.append((ok, f"Source_Prices!{sheet.cell(row, column).coordinate} 来源值错误"))
    else:
        source_checks.append((False, "缺少 Source_Prices"))

    native.extend([
        (sum(len(sheet.tables) for sheet in workbook.worksheets) >= 3, "Excel 原生表少于 3"),
        (sum(len(sheet._charts) for sheet in workbook.worksheets) >= 1, "Excel 缺少原生行业图表"),
        (all(name in workbook.sheetnames and workbook[name].freeze_panes == "A2" for name in ("Holdings", "Summary", "Source_Prices")), "Excel 未全部冻结表头"),
        (all(name in workbook.sheetnames and bool(workbook[name].auto_filter.ref) for name in ("Holdings", "Summary", "Source_Prices")), "Excel 未全部启用筛选"),
        (sum(len(sheet.conditional_formatting) for sheet in workbook.worksheets) >= 3, "Excel 条件格式不足"),
    ])
    return {
        "identity": _ratio(identity), "boundary": _ratio(boundary), "calculations": _ratio(calculations),
        "policy": _ratio(policy), "summary": _ratio(summary_checks), "source": _ratio(source_checks),
        "native": native, "symbols": delivered_symbols, "formula_count": formula_count,
        "formula_total": formula_total, "workbook": workbook,
    }


def _word(path: Path, truth, summary):
    document = Document(path)
    text = V1.word_text(document)
    blob = _lower(text)
    headings = [_norm(p.text) for p in document.paragraphs if str(p.style.name).startswith("Heading")]
    holdings: list[tuple[bool, str]] = [
        ("Portfolio Analysis Report" in text, "Word 标题错误"),
        *[(heading in headings, f"Word 缺章节 {heading}") for heading in ("Holdings Overview", "Performance Analysis", "Policy Review", "Recommendations")],
        (len(document.tables) >= 3, "Word 原生表少于 3"),
    ]
    row_texts = [" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows[1:]]
    records = {
        str(item["symbol"]): max((row for row in row_texts if str(item["symbol"]) in row), key=len, default="")
        for item in truth
    }
    for expected in truth:
        row = records[str(expected["symbol"])]
        tokens = [
            expected["symbol"], expected["company"], expected["sector"], f"{int(expected['quantity']):,}",
            f"{Decimal(expected['current_price']):,.4f}", V1.money(Decimal(expected["current_value"])),
            f"{Decimal(expected['return_pct']):.2%}", f"{Decimal(expected['weight_pct']):.2%}",
        ]
        for token in tokens:
            holdings.append((_lower(token) in _lower(row), f"Word {expected['symbol']} 缺 {token}"))

    claims = []
    for token in (
        "192,071.02 RUB", "MTSS.ME", "1.96%", "SBER.ME", "-5.41%", "Energy",
        "78,997.09 RUB", "41.13%", "5 of 5", "does not define a maximum position or sector-weight limit",
    ):
        claims.append((_lower(token) in blob, f"Word 缺结论 {token}"))

    recommendation_paragraphs = []
    active = False
    for paragraph in document.paragraphs:
        value = _norm(paragraph.text)
        if value == "Recommendations":
            active = True
            continue
        if active and str(paragraph.style.name).startswith("Heading"):
            break
        if active and value:
            recommendation_paragraphs.append(value)
    recommendation_blob = _lower("\n".join(recommendation_paragraphs))
    recommendations = [
        (len(recommendation_paragraphs) >= 2, f"建议少于两条：{len(recommendation_paragraphs)}"),
        ("sber.me" in recommendation_blob and ("weak" in recommendation_blob or "-5.41%" in recommendation_blob), "建议未锚定最弱持仓"),
        ("energy" in recommendation_blob and "41.13%" in recommendation_blob, "建议未锚定 Energy 暴露"),
        ("frozen" in blob or "snapshot" in blob or "observed" in blob or "evidence" in blob, "未声明冻结证据边界"),
        ("cause" in blob and ("do not" in blob or "does not" in blob or "without" in blob), "未声明同期变化不证明因果"),
    ]
    return _ratio(holdings), _ratio(claims), _ratio(recommendations), text, document, headings


def _email(path: Path, artifact: Path):
    message = _message(path)
    body = _mail_body(message)
    content = [
        (_lower(message.get("From")) == "analyst@fund.example.com", "邮件发件人错误"),
        (_lower(message.get("To")) == "investment-committee@fund.example.com", "邮件收件人错误"),
        (_norm(message.get("Subject")) == "Monthly Portfolio Analysis Report", "邮件主题错误"),
    ]
    for token in ("2026-05-26", "192,071.02 RUB", "MTSS.ME", "1.96%", "SBER.ME", "-5.41%"):
        content.append((_lower(token) in _lower(body), f"邮件正文缺 {token}"))

    attachments: dict[str, tuple[bytes, str]] = {}
    for part in message.iter_attachments():
        name = _norm(part.get_filename())
        if name:
            attachments[name] = (part.get_payload(decode=True) or b"", part.get_content_type())
    attachment_checks = [(set(attachments) == {XLSX, DOCX}, f"邮件附件集合错误：{sorted(attachments)}")]
    expected_mime = {
        XLSX: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        DOCX: "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }
    for name in (XLSX, DOCX):
        payload, mime = attachments.get(name, (b"", ""))
        attachment_checks.extend([
            (bool(payload), f"附件 {name} 缺失/为空"),
            (mime == expected_mime[name], f"附件 {name} MIME 错误"),
            ((artifact / name).is_file() and payload == (artifact / name).read_bytes(), f"附件 {name} 与目录文件不一致"),
        ])
    return _ratio(content), _ratio(attachment_checks), body, message, attachments


def _render_findings(artifact: Path):
    render = V1.render_office(artifact)
    return [
        (not render["errors"], f"LibreOffice 渲染错误：{render['errors']}"),
        (5 <= render["xlsx_pages"] <= 8, f"Excel 渲染页数异常：{render['xlsx_pages']}"),
        (2 <= render["docx_pages"] <= 4, f"Word 渲染页数异常：{render['docx_pages']}"),
    ]


def _professional(artifact: Path, workbook, document, native):
    findings = [*native, (len(document.tables) >= 3, "Word 原生表少于 3")]
    package_info = {}
    for name in (XLSX, DOCX):
        info = V1.inspect_package(artifact / name)
        package_info[name] = info
        findings.extend([
            (not info["broken"] and not info["missing_targets"], f"{name} OOXML 关系损坏"),
            (not info["macros"], f"{name} 含宏"),
            (not info["external"], f"{name} 含外部关系"),
        ])
    findings.append((not package_info[DOCX]["media"], "Word 持仓表被图片替代"))
    findings.extend(_render_findings(artifact))
    return _ratio(findings)


def _identity_penalty(symbols, truth):
    expected = {str(item["symbol"]) for item in truth}
    delivered = [symbol for symbol in symbols if symbol]
    missing = expected - set(delivered)
    extra = set(delivered) - expected
    duplicate = len(delivered) - len(set(delivered))
    affected = len(missing) + len(extra) + max(0, duplicate)
    severity = min(1.0, affected / max(1, len(expected)))
    return severity, ([] if not affected else [f"代码缺失={sorted(missing)}，额外={sorted(extra)}，重复={duplicate}"]), len(expected)


def _parse_money(text: str):
    match = re.search(r"([0-9][0-9, ]+(?:\.[0-9]+)?)\s*RUB", text, flags=re.IGNORECASE)
    if not match:
        return None
    try:
        return float(match.group(1).replace(",", "").replace(" ", ""))
    except ValueError:
        return None


def _explicit_conflicts(text: str, *, carrier: str):
    normalized = _norm(text)
    conflicts = []
    for pattern in (r"total portfolio value is\s+([0-9][0-9, ]+(?:\.[0-9]+)?)\s*RUB", r"portfolio value as of [^ ]+ is\s+([0-9][0-9, ]+(?:\.[0-9]+)?)\s*RUB"):
        match = re.search(pattern, normalized, flags=re.IGNORECASE)
        if match and not _approx(match.group(1).replace(",", "").replace(" ", ""), 192071.0245, 0.02):
            conflicts.append(f"{carrier} 总市值冲突")
    patterns = (
        (r"([A-Z]{4}\.ME)\s+is the best", "MTSS.ME", "最佳代码"),
        (r"best 30-day performer:\s*([A-Z]{4}\.ME)", "MTSS.ME", "最佳代码"),
        (r"([A-Z]{4}\.ME)\s+is the worst", "SBER.ME", "最差代码"),
        (r"worst 30-day performer:\s*([A-Z]{4}\.ME)", "SBER.ME", "最差代码"),
        (r"([A-Za-z ]+)\s+is the largest sector exposure", "Energy", "最大行业"),
    )
    for pattern, expected, label in patterns:
        match = re.search(pattern, normalized, flags=re.IGNORECASE)
        if match and _norm(match.group(1)).lower() != expected.lower():
            conflicts.append(f"{carrier} {label}冲突")
    return conflicts


def _cross_penalty(word_text: str, mail_body: str):
    issues = [*_explicit_conflicts(word_text, carrier="Word"), *_explicit_conflicts(mail_body, carrier="EML")]
    severity = min(1.0, 0.35 * len(issues))
    return severity, issues, 1.0


def _email_penalty(message, attachments, artifact: Path):
    issues: list[tuple[float, str]] = []
    if _lower(message.get("To")) != "investment-committee@fund.example.com":
        issues.append((0.8, "邮件错投"))
    if _norm(message.get("Subject")) != "Monthly Portfolio Analysis Report":
        issues.append((0.4, "邮件主题错误"))
    if set(attachments) != {XLSX, DOCX}:
        issues.append((1.0, "邮件附件缺失、额外或命名错误"))
    else:
        stale = [name for name in (XLSX, DOCX) if attachments[name][0] != (artifact / name).read_bytes()]
        if stale:
            issues.append((1.0, f"邮件附件陈旧：{stale}"))
    severity = max((value for value, _ in issues), default=0.0)
    return severity, [label for _, label in issues], 1.0


def _safe_penalty(artifact: Path, formula_count: int, formula_total: int):
    coverage = formula_count / max(1, formula_total)
    issues: list[tuple[float, str]] = []
    if coverage < 0.5:
        issues.append((min(1.0, (0.5 - coverage) / 0.5), f"原生公式覆盖率仅 {coverage:.1%}"))
    extras = sorted(path.name for path in artifact.iterdir() if path.name not in {*REQUIRED, "__pycache__", "solve.py", "solve.sh"})
    if extras:
        issues.append((0.4, f"额外输出：{extras}"))
    suspicious = V1.suspicious_content(artifact)
    if suspicious:
        issues.append((0.8, f"绝对路径/Golden/Judge 泄漏：{suspicious}"))
    unsafe = []
    for name in (XLSX, DOCX):
        info = V1.inspect_package(artifact / name)
        if info["macros"] or info["external"]:
            unsafe.append(name)
    if unsafe:
        issues.append((1.0, f"宏或危险外链：{unsafe}"))
    severity = max((value for value, _ in issues), default=0.0)
    return severity, [label for _, label in issues], 1.0


def _unsupported_penalty(text: str):
    findings = V1.unsupported_claims(text)
    severity = min(1.0, len(findings) / 2)
    return severity, findings, 1.0


def evaluate(artifact: Path):
    core_ok, core_errors, hashes, truth, source_rows, summary = _core(artifact)
    if not core_ok:
        return _fatal(core_errors, hashes)

    workbook = _workbook(artifact / XLSX, truth, source_rows, summary)
    word_holdings, word_claims, recommendations, word_text, document, _headings = _word(artifact / DOCX, truth, summary)
    mail_content, mail_attachments, mail_body, message, attachments = _email(artifact / EML, artifact)
    professional = _professional(artifact, workbook["workbook"], document, workbook["native"])

    groups = {
        "R002": workbook["identity"], "R003": workbook["boundary"], "R004": workbook["calculations"],
        "R005": workbook["policy"], "R006": workbook["summary"], "R007": workbook["source"],
        "R008": word_holdings, "R009": word_claims, "R010": recommendations,
        "R011": mail_content, "R012": mail_attachments, "R013": professional,
    }
    results = [_criterion("R001", 1.0, "三个核心交付可解析且冻结输入哈希正确", 3, 3)]
    scores = {"R001": 1.0}
    for key in [f"R{index:03d}" for index in range(2, 14)]:
        score, failures, earned, total = groups[key]
        evidence = "全部任务专属原子检查通过" if not failures else "未通过：" + "；".join(str(item) for item in failures)
        results.append(_criterion(key, score, evidence, earned, total))
        scores[key] = score

    penalties = {
        "R014": _identity_penalty(workbook["symbols"], truth),
        "R015": _cross_penalty(word_text, mail_body),
        "R016": _email_penalty(message, attachments, artifact),
        "R017": _safe_penalty(artifact, workbook["formula_count"], workbook["formula_total"]),
        "R018": _unsupported_penalty("\n".join((word_text, mail_body))),
    }
    for key, (severity, labels, total) in penalties.items():
        results.append(_criterion(key, severity, "未触发任务专属严重边界" if not labels else "；".join(labels), severity * total, total))
        scores[key] = severity

    positive_total = sum(POSITIVE_WEIGHTS.values())
    positive = sum(POSITIVE_WEIGHTS[key] * scores[key] for key in POSITIVE_WEIGHTS)
    penalty = sum(abs(PENALTY_WEIGHTS[key]) * scores[key] for key in PENALTY_WEIGHTS)
    reward = round(max(0.0, min(1.0, (positive - penalty) / positive_total)), 6)
    workbook["workbook"].close()
    return {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes, "criterion_results": results, "reward": reward,
        "passed": reward >= 0.8, "hurdle_failures": [],
        "penalties": [{"id": key, "amount": round(abs(PENALTY_WEIGHTS[key]) * scores[key] / positive_total, 6)}
                      for key in PENALTY_WEIGHTS if scores[key] > 0],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    parser.add_argument("--scoring-version", choices=("v1", "v2"), default="v2")
    args = parser.parse_args()
    try:
        result = V1.evaluate(args.artifact_dir) if args.scoring_version == "v1" else evaluate(args.artifact_dir)
    except Exception as exc:  # noqa: BLE001
        result = _fatal([f"Judge 异常：{type(exc).__name__}: {exc}"], [])
    output = json.dumps(result, ensure_ascii=False, indent=2)
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(output + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{float(result['reward']):.6f}\n", encoding="utf-8")
    print(output)


if __name__ == "__main__":
    main()
