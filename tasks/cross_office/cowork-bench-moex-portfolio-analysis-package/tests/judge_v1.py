#!/usr/bin/env python3
"""MOEX 投资组合跨 Office 任务的来源重算、跨文件一致性与安全 Judge。"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from email import policy
from email.parser import BytesParser
from pathlib import Path

from docx import Document
from lxml import etree
from openpyxl import load_workbook


TASK_ID = "cowork-evo/cowork-bench-moex-portfolio-analysis-package"
RUBRIC_VERSION = "cowork-bench-moex-portfolio-analysis-v0.1"
JUDGE_VERSION = "cowork-bench-moex-portfolio-analysis-deterministic-v0.1"
XLSX_NAME = "Portfolio_Analysis.xlsx"
DOCX_NAME = "Investment_Memo.docx"
EML_NAME = "Monthly_Portfolio_Analysis.eml"
REQUIRED_FILES = [XLSX_NAME, DOCX_NAME, EML_NAME]
SOURCE_HASHES = {
    "Portfolio.xlsx": "a51b31211a425c99f07633eba139082f5737f086f273aedd575edeb8bdfca06d",
    "Investment_Policy.pdf": "475b991765883c7030d8ec8e191a209781dc51763a955c865d4beaba9538b02d",
    "moex_market_snapshot.csv": "b760d32e925c23413f7d58395f80e778313692923615523c13de9912abbf73d4",
}
HOLDING_HEADERS = [
    "Symbol", "Company", "Sector", "Quantity", "Baseline_Date", "Baseline_Close", "Current_Date",
    "Current_Price", "Current_Value", "Return_30d_Pct", "Weight_Pct", "Market_Cap_RUB",
    "Market_Cap_Category", "Currency", "Exchange", "Policy_Status",
]
SOURCE_HEADERS = ["Symbol", "Company", "Sector", "Currency", "Exchange", "Market_Cap", "Date", "Close"]
SUMMARY_LABELS = [
    "Total_Portfolio_Value", "Best_Performer", "Worst_Performer", "Financial Services", "Energy",
    "Consumer Defensive", "Communication Services", "Largest_Sector", "Largest_Sector_Weight",
    "Policy_Compliant_Positions",
]
ALLOWED_SECTORS = {
    "Communication Services", "Consumer Defensive", "Financial Services", "Healthcare", "Energy"
}
SUSPICIOUS_MARKERS = ["solution/golden", "judge.py", "expectations.json", "judge-result", "/root/", "/home/", "c:\\"]
US_TICKERS = ["GOOGL", "AMZN", "JPM", "JNJ", "XOM"]

SPECS = {
    "R001": ("完整交付物与冻结输入", 5, "hurdle", 0.0),
    "R002": ("五个持仓与发行人字段", 7, "hurdle", 0.3),
    "R003": ("冻结行情边界与价格", 8, "hurdle", 0.35),
    "R004": ("持仓计算公式与真值", 12, "hurdle", 0.35),
    "R005": ("市值分类与政策公式", 7, "hurdle", 0.5),
    "R006": ("公式化汇总与行业暴露", 9, "hurdle", 0.45),
    "R007": ("一百一十五行来源追溯", 8, "hurdle", 0.5),
    "R008": ("Word 持仓表与可编辑结构", 9, "hurdle", 0.5),
    "R009": ("Word 表现与政策结论", 7, "hurdle", 0.55),
    "R010": ("有依据且克制的建议", 5, "hurdle", 0.6),
    "R011": ("邮件投递头与正文", 6, "hurdle", 0.5),
    "R012": ("邮件附件字节一致", 8, "hurdle", 0.4),
    "R013": ("原生结构与可渲染性", 9, "positive", None),
    "R014": ("无持仓、行情或计算冲突", -30, "penalty", None),
    "R015": ("三份交付物事实一致", -25, "penalty", None),
    "R016": ("无错投、缺附件或陈旧附件", -30, "penalty", None),
    "R017": ("无损坏、静态化、外链或调试泄漏", -25, "penalty", None),
    "R018": ("无上游错误美股数据或无来源断言", -15, "penalty", None),
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def normalize(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\u00a0", " ").replace("\v", " ")).strip()


def approx(left: object, right: object, tolerance: float = 1e-6) -> bool:
    try:
        return abs(float(left) - float(right)) <= tolerance
    except (TypeError, ValueError):
        return False


def money(value: Decimal) -> str:
    return f"{value.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP):,.2f}"


def source_path(name: str) -> Path | None:
    task_root = Path(__file__).resolve().parents[1]
    candidates = [
        Path("/app/input_files") / name,
        task_root / "data" / "input_files" / name,
        task_root / "environment" / "input_files" / name,
    ]
    return next((path for path in candidates if path.is_file()), None)


def load_truth() -> tuple[list[dict[str, object]], list[dict[str, object]], dict[str, str], dict[str, object]]:
    errors: dict[str, str] = {}
    paths: dict[str, Path] = {}
    for name, expected in SOURCE_HASHES.items():
        path = source_path(name)
        if path is None:
            errors[name] = "missing"
            continue
        paths[name] = path
        actual = sha256(path)
        if actual != expected:
            errors[name] = f"sha256={actual}"
    if errors:
        return [], [], errors, {}

    workbook = load_workbook(paths["Portfolio.xlsx"], data_only=False, read_only=True)
    sheet = workbook["Holdings"]
    portfolio = [
        {"symbol": str(row[0]), "company": str(row[1]), "quantity": int(row[2])}
        for row in sheet.iter_rows(min_row=2, values_only=True) if any(value is not None for value in row)
    ]
    workbook.close()

    market_rows: list[dict[str, object]] = []
    with paths["moex_market_snapshot.csv"].open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames != SOURCE_HEADERS:
            errors["moex_market_snapshot.csv"] = f"headers={reader.fieldnames}"
            return [], [], errors, {}
        for row in reader:
            market_rows.append(
                {
                    "symbol": row["Symbol"], "company": row["Company"], "sector": row["Sector"],
                    "currency": row["Currency"], "exchange": row["Exchange"],
                    "market_cap": int(row["Market_Cap"]), "date": date.fromisoformat(row["Date"]),
                    "close": Decimal(row["Close"]),
                }
            )
    by_symbol: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in market_rows:
        by_symbol[str(row["symbol"])].append(row)

    as_of = date(2026, 5, 26)
    cutoff = as_of - timedelta(days=30)
    analyzed: list[dict[str, object]] = []
    for holding in portfolio:
        history = sorted(by_symbol[str(holding["symbol"])], key=lambda item: item["date"])
        latest = max(history, key=lambda item: item["date"])
        baseline = max((item for item in history if item["date"] <= cutoff), key=lambda item: item["date"])
        facts = {
            (item["company"], item["sector"], item["currency"], item["exchange"], item["market_cap"])
            for item in history
        }
        company, sector, currency, exchange, market_cap = facts.pop()
        quantity = int(holding["quantity"])
        value = Decimal(quantity) * Decimal(latest["close"])
        return_pct = Decimal(latest["close"]) / Decimal(baseline["close"]) - Decimal("1")
        analyzed.append(
            {
                **holding, "company": company, "sector": sector, "currency": currency, "exchange": exchange,
                "market_cap": market_cap, "baseline_date": baseline["date"], "baseline_close": baseline["close"],
                "current_date": latest["date"], "current_price": latest["close"], "current_value": value,
                "return_pct": return_pct,
            }
        )
    total = sum((Decimal(item["current_value"]) for item in analyzed), Decimal("0"))
    sectors: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    for item in analyzed:
        item["weight_pct"] = Decimal(item["current_value"]) / total
        item["market_cap_category"] = "Large Cap" if int(item["market_cap"]) > 100_000_000_000 else "Below Large Cap"
        item["policy_status"] = (
            "Compliant" if item["sector"] in ALLOWED_SECTORS and item["currency"] == "RUB"
            and item["exchange"] == "MCX" and item["market_cap_category"] == "Large Cap" else "Review"
        )
        sectors[str(item["sector"])] += Decimal(item["current_value"])
    best = max(analyzed, key=lambda item: Decimal(item["return_pct"]))
    worst = min(analyzed, key=lambda item: Decimal(item["return_pct"]))
    largest_sector = max(sectors, key=sectors.get)
    summary = {
        "total": total, "best": best, "worst": worst, "sectors": sectors, "largest_sector": largest_sector,
        "largest_sector_weight": sectors[largest_sector] / total,
        "compliant": sum(item["policy_status"] == "Compliant" for item in analyzed),
    }
    return analyzed, market_rows, errors, summary


def iso_date(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = normalize(value)
    try:
        return date.fromisoformat(text[:10]).isoformat()
    except ValueError:
        return None


def inspect_package(path: Path) -> dict[str, object]:
    result: dict[str, object] = {"broken": None, "missing_targets": [], "macros": [], "external": [], "media": []}
    if not path.is_file():
        result["broken"] = "missing"
        return result
    try:
        with zipfile.ZipFile(path) as package:
            names = set(package.namelist())
            result["macros"] = sorted(name for name in names if "vbaProject" in name or name.endswith(".bin"))
            result["media"] = sorted(name for name in names if "/media/" in name)
            for rel_name in [name for name in names if name.endswith(".rels")]:
                root = etree.fromstring(package.read(rel_name))
                for rel in root:
                    target = rel.get("Target", "")
                    if rel.get("TargetMode") == "External":
                        result["external"].append(target)
                        continue
                    if not target or target.startswith("#"):
                        continue
                    if target.startswith("/"):
                        resolved = target.lstrip("/")
                    else:
                        base = Path(rel_name).parent
                        if base.name == "_rels":
                            base = base.parent
                        resolved = str((base / target).as_posix())
                        while "/../" in resolved:
                            resolved = re.sub(r"[^/]+/\.\./", "", resolved, count=1)
                    if resolved not in names:
                        result["missing_targets"].append({"rels": rel_name, "target": target, "resolved": resolved})
    except Exception as exc:
        result["broken"] = f"{type(exc).__name__}: {exc}"
    return result


def recalculate_candidate(path: Path) -> tuple[Path | None, str]:
    temp_root = Path(tempfile.mkdtemp(prefix="moex-portfolio-judge-recalc-"))
    input_root = temp_root / "input"
    output_root = temp_root / "output"
    profile = temp_root / "profile"
    input_root.mkdir()
    output_root.mkdir()
    profile.mkdir()
    source = input_root / path.name
    shutil.copy2(path, source)
    command = [
        "soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "xlsx",
        "--outdir", str(output_root), str(source),
    ]
    try:
        proc = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300)
    except Exception as exc:
        return None, f"{type(exc).__name__}: {exc}"
    converted = output_root / path.name
    if proc.returncode != 0 or not converted.is_file():
        return None, proc.stdout[-1000:]
    return converted, proc.stdout[-1000:]


def parse_workbook(path: Path, truth: list[dict[str, object]], source_rows: list[dict[str, object]], summary: dict[str, object]) -> dict[str, object]:
    result: dict[str, object] = {
        "readable": False, "sheetnames": [], "identity_errors": [], "boundary_errors": [], "calculation_errors": [],
        "policy_errors": [], "summary_errors": [], "source_errors": [], "formula_errors": [], "native": {},
    }
    try:
        formulas = load_workbook(path, data_only=False)
    except Exception as exc:
        result["error"] = f"{type(exc).__name__}: {exc}"
        return result
    recalculated, log = recalculate_candidate(path)
    result["recalc_log"] = log
    if recalculated is None:
        return result
    try:
        values = load_workbook(recalculated, data_only=True)
    except Exception as exc:
        result["error"] = f"recalc load: {type(exc).__name__}: {exc}"
        return result
    result["readable"] = True
    result["sheetnames"] = formulas.sheetnames
    if formulas.sheetnames != ["Holdings", "Summary", "Source_Prices"]:
        result["identity_errors"].append(f"sheetnames={formulas.sheetnames}")
        return result

    holdings_f = formulas["Holdings"]
    holdings_v = values["Holdings"]
    headers = [normalize(cell.value) for cell in holdings_f[1]]
    if headers != HOLDING_HEADERS:
        result["identity_errors"].append(f"headers={headers}")
    if holdings_f.max_row != 6 or holdings_f.max_column != 16:
        result["identity_errors"].append(f"shape={holdings_f.max_row}x{holdings_f.max_column}")
    for row_index, expected in enumerate(truth, start=2):
        identity = [expected["symbol"], expected["company"], expected["sector"], expected["quantity"], expected["market_cap"], expected["currency"], expected["exchange"]]
        for column, expected_value in zip((1, 2, 3, 4, 12, 14, 15), identity, strict=True):
            got = holdings_v.cell(row_index, column).value
            if normalize(got) != normalize(expected_value):
                result["identity_errors"].append(f"{holdings_v.cell(row_index, column).coordinate}={got!r}")
        boundaries = [
            (5, expected["baseline_date"].isoformat(), "date"), (6, expected["baseline_close"], "number"),
            (7, expected["current_date"].isoformat(), "date"), (8, expected["current_price"], "number"),
        ]
        for column, expected_value, kind in boundaries:
            got = holdings_v.cell(row_index, column).value
            valid = iso_date(got) == expected_value if kind == "date" else approx(got, expected_value, 1e-7)
            if not valid:
                result["boundary_errors"].append(f"{holdings_v.cell(row_index, column).coordinate}={got!r}")
        for column in (9, 10, 11, 13, 16):
            if holdings_f.cell(row_index, column).data_type != "f":
                result["formula_errors"].append(f"Holdings!{holdings_f.cell(row_index, column).coordinate}:not_formula")
        formula_i = normalize(holdings_f.cell(row_index, 9).value).upper().replace("$", "")
        formula_j = normalize(holdings_f.cell(row_index, 10).value).upper().replace("$", "")
        formula_k = normalize(holdings_f.cell(row_index, 11).value).upper().replace("$", "")
        formula_m = normalize(holdings_f.cell(row_index, 13).value).upper().replace("$", "")
        formula_p = normalize(holdings_f.cell(row_index, 16).value).upper().replace("$", "")
        if f"D{row_index}*H{row_index}" not in formula_i:
            result["formula_errors"].append(f"I{row_index}:wrong_source")
        if f"H{row_index}/F{row_index}-1" not in formula_j:
            result["formula_errors"].append(f"J{row_index}:wrong_source")
        if f"I{row_index}/SUM(I2:I6)" not in formula_k:
            result["formula_errors"].append(f"K{row_index}:wrong_source")
        if f"L{row_index}>100000000000" not in formula_m:
            result["formula_errors"].append(f"M{row_index}:wrong_threshold")
        if not all(f"{column}{row_index}" in formula_p for column in ("C", "N", "O", "M")):
            result["formula_errors"].append(f"P{row_index}:wrong_source")
        calculated = [holdings_v.cell(row_index, column).value for column in (9, 10, 11)]
        expected_numbers = [expected["current_value"], expected["return_pct"], expected["weight_pct"]]
        for column, got, wanted in zip((9, 10, 11), calculated, expected_numbers, strict=True):
            if not approx(got, wanted, 1e-7):
                result["calculation_errors"].append(f"{holdings_v.cell(row_index, column).coordinate}={got!r} expected={wanted}")
        if normalize(holdings_v.cell(row_index, 13).value) != expected["market_cap_category"]:
            result["policy_errors"].append(f"M{row_index}={holdings_v.cell(row_index, 13).value!r}")
        if normalize(holdings_v.cell(row_index, 16).value) != expected["policy_status"]:
            result["policy_errors"].append(f"P{row_index}={holdings_v.cell(row_index, 16).value!r}")

    summary_f = formulas["Summary"]
    summary_v = values["Summary"]
    labels = [normalize(summary_v.cell(row, 1).value) for row in range(2, 12)]
    if labels != SUMMARY_LABELS:
        result["summary_errors"].append(f"labels={labels}")
    sector_values = summary["sectors"]
    expected_values = [
        summary["total"], summary["best"]["symbol"], summary["worst"]["symbol"],
        sector_values["Financial Services"], sector_values["Energy"], sector_values["Consumer Defensive"],
        sector_values["Communication Services"], summary["largest_sector"], summary["largest_sector_weight"], summary["compliant"],
    ]
    for row, expected_value in enumerate(expected_values, start=2):
        cell_f = summary_f.cell(row, 2)
        got = summary_v.cell(row, 2).value
        if cell_f.data_type != "f":
            result["formula_errors"].append(f"Summary!B{row}:not_formula")
        if isinstance(expected_value, (Decimal, int, float)):
            if not approx(got, expected_value, 1e-7):
                result["summary_errors"].append(f"B{row}={got!r} expected={expected_value}")
        elif normalize(got) != normalize(expected_value):
            result["summary_errors"].append(f"B{row}={got!r} expected={expected_value!r}")

    source = formulas["Source_Prices"]
    if [normalize(cell.value) for cell in source[1]] != SOURCE_HEADERS:
        result["source_errors"].append("wrong_headers")
    if source.max_row != 116 or source.max_column != 8:
        result["source_errors"].append(f"shape={source.max_row}x{source.max_column}")
    for row_index, expected in enumerate(source_rows, start=2):
        got = [source.cell(row_index, column).value for column in range(1, 9)]
        wanted = [
            expected["symbol"], expected["company"], expected["sector"], expected["currency"], expected["exchange"],
            expected["market_cap"], expected["date"].isoformat(), expected["close"],
        ]
        for column, (actual, target) in enumerate(zip(got, wanted, strict=True), start=1):
            if column == 7:
                valid = iso_date(actual) == target
            elif column in (6, 8):
                valid = approx(actual, target, 1e-7)
            else:
                valid = normalize(actual) == normalize(target)
            if not valid:
                result["source_errors"].append(f"{source.cell(row_index, column).coordinate}={actual!r}")
                if len(result["source_errors"]) >= 25:
                    break
        if len(result["source_errors"]) >= 25:
            break

    formula_error_cells = []
    for sheet in values.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and re.match(r"^#(?:NAME|VALUE|REF|DIV/0|N/A)", cell.value, flags=re.IGNORECASE):
                    formula_error_cells.append(f"{sheet.title}!{cell.coordinate}={cell.value}")
    result["formula_errors"].extend(formula_error_cells)
    result["native"] = {
        "tables": sum(len(sheet.tables) for sheet in formulas.worksheets),
        "charts": sum(len(sheet._charts) for sheet in formulas.worksheets),
        "freeze": all(sheet.freeze_panes == "A2" for sheet in formulas.worksheets),
        "filters": all(bool(sheet.auto_filter.ref) for sheet in formulas.worksheets),
        "conditional_formats": sum(len(sheet.conditional_formatting) for sheet in formulas.worksheets),
    }
    return result


def word_text(document: Document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def unsupported_claims(text: str) -> list[str]:
    findings: list[str] = []
    for ticker in US_TICKERS:
        if re.search(rf"\b{re.escape(ticker)}(?:\.ME)?\b", text, flags=re.IGNORECASE):
            findings.append(f"us_ticker:{ticker}")
    if re.search(r"\bUSD\b|\$\s*87[, ]?030", text, flags=re.IGNORECASE):
        findings.append("upstream_usd_data")
    if re.search(r"(?:policy|sector|position).{0,45}(?:limit|maximum|cap).{0,20}\b\d+(?:\.\d+)?\s*%", text, flags=re.IGNORECASE):
        findings.append("invented_policy_threshold")
    if re.search(r"driven by favorable|because of strong management|automatic (?:buy|sell) recommendation|guaranteed return", text, flags=re.IGNORECASE):
        findings.append("unsupported_causal_or_trading_claim")
    return sorted(set(findings))


def parse_word(path: Path, truth: list[dict[str, object]], summary: dict[str, object]) -> dict[str, object]:
    result: dict[str, object] = {
        "readable": False, "holding_errors": [], "claim_errors": [], "recommendation_errors": [],
        "unsupported": [], "tables": 0, "headings": [], "recommendation_count": 0,
    }
    try:
        document = Document(path)
    except Exception as exc:
        result["error"] = f"{type(exc).__name__}: {exc}"
        return result
    result["readable"] = True
    result["tables"] = len(document.tables)
    result["headings"] = [normalize(p.text) for p in document.paragraphs if str(p.style.name).startswith("Heading")]
    text = word_text(document)
    normalized = normalize(text)
    if "Portfolio Analysis Report" not in normalized:
        result["claim_errors"].append("wrong_title")
    for heading in ("Holdings Overview", "Performance Analysis", "Policy Review", "Recommendations"):
        if heading not in result["headings"]:
            result["claim_errors"].append(f"missing_heading:{heading}")
    for expected in truth:
        tokens = [
            expected["symbol"], expected["company"], expected["sector"], f"{int(expected['quantity']):,}",
            f"{Decimal(expected['current_price']):,.4f}", money(Decimal(expected["current_value"])),
            f"{Decimal(expected['return_pct']):.2%}", f"{Decimal(expected['weight_pct']):.2%}",
        ]
        for token in tokens:
            if normalize(token).lower() not in normalized.lower():
                result["holding_errors"].append(f"{expected['symbol']}:missing:{token}")
    required_claims = [
        "192,071.02 RUB", "MTSS.ME", "1.96%", "SBER.ME", "-5.41%", "Energy", "78,997.09 RUB", "41.13%", "5 of 5",
        "does not define a maximum position or sector-weight limit",
    ]
    for token in required_claims:
        if token.lower() not in normalized.lower():
            result["claim_errors"].append(f"missing:{token}")
    in_recommendations = False
    for paragraph in document.paragraphs:
        paragraph_text = normalize(paragraph.text)
        if paragraph_text == "Recommendations":
            in_recommendations = True
            continue
        if in_recommendations and str(paragraph.style.name).startswith("Heading"):
            break
        if in_recommendations and paragraph_text:
            result["recommendation_count"] += 1
    if result["recommendation_count"] < 2:
        result["recommendation_errors"].append(f"count={result['recommendation_count']}")
    if not re.search(r"frozen|snapshot|observed|evidence", normalized, flags=re.IGNORECASE):
        result["recommendation_errors"].append("no_evidence_boundary")
    result["unsupported"] = unsupported_claims(normalized)
    return result


def parse_email(path: Path, artifact_dir: Path) -> dict[str, object]:
    result: dict[str, object] = {
        "readable": False, "header_errors": [], "body_errors": [], "attachment_errors": [],
        "unsupported": [], "attachments": [], "body": "",
    }
    try:
        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
    except Exception as exc:
        result["error"] = f"{type(exc).__name__}: {exc}"
        return result
    result["readable"] = True
    expected_headers = {
        "From": "analyst@fund.example.com", "To": "investment-committee@fund.example.com",
        "Subject": "Monthly Portfolio Analysis Report",
    }
    for name, expected in expected_headers.items():
        if normalize(message.get(name, "")).lower() != expected.lower():
            result["header_errors"].append(f"{name}={message.get(name)!r}")
    try:
        body_part = message.get_body(preferencelist=("plain",)) if message.is_multipart() else message
        body = body_part.get_content() if body_part else ""
    except Exception:
        body = ""
    result["body"] = normalize(body)
    for token in ("2026-05-26", "192,071.02 RUB", "MTSS.ME", "1.96%", "SBER.ME", "-5.41%"):
        if token.lower() not in result["body"].lower():
            result["body_errors"].append(f"missing:{token}")
    expected_mime = {
        XLSX_NAME: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        DOCX_NAME: "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }
    attachments = list(message.iter_attachments())
    if len(attachments) != 2:
        result["attachment_errors"].append(f"count={len(attachments)}")
    for part in attachments:
        filename = normalize(part.get_filename())
        payload = part.get_payload(decode=True) or b""
        details = {"filename": filename, "mime": part.get_content_type(), "sha256": hashlib.sha256(payload).hexdigest()}
        result["attachments"].append(details)
        if filename not in expected_mime:
            result["attachment_errors"].append(f"unexpected:{filename}")
            continue
        if part.get_content_type() != expected_mime[filename]:
            result["attachment_errors"].append(f"mime:{filename}={part.get_content_type()}")
        delivered = artifact_dir / filename
        if not delivered.is_file() or hashlib.sha256(payload).hexdigest() != sha256(delivered):
            result["attachment_errors"].append(f"bytes:{filename}")
    if sorted(item["filename"] for item in result["attachments"]) != sorted(expected_mime):
        result["attachment_errors"].append("wrong_attachment_names")
    result["unsupported"] = unsupported_claims(result["body"])
    return result


def render_office(artifact_dir: Path) -> dict[str, object]:
    result: dict[str, object] = {"xlsx_pages": 0, "docx_pages": 0, "errors": [], "logs": []}
    with tempfile.TemporaryDirectory(prefix="moex-portfolio-judge-render-") as tmp:
        root = Path(tmp)
        output = root / "output"
        profile = root / "profile"
        output.mkdir()
        profile.mkdir()
        for name in (XLSX_NAME, DOCX_NAME):
            source = artifact_dir / name
            if not source.is_file():
                result["errors"].append(f"missing:{name}")
                continue
            command = [
                "soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf",
                "--outdir", str(output), str(source),
            ]
            try:
                proc = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300)
            except Exception as exc:
                result["errors"].append(f"{name}:{type(exc).__name__}:{exc}")
                continue
            result["logs"].append(proc.stdout[-1000:])
            pdf = output / f"{Path(name).stem}.pdf"
            if proc.returncode != 0 or not pdf.is_file():
                result["errors"].append(f"{name}:conversion_failed")
                continue
            info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=60)
            match = re.search(r"^Pages:\s+(\d+)", info.stdout, flags=re.MULTILINE)
            pages = int(match.group(1)) if match else 0
            extracted = subprocess.run(["pdftotext", str(pdf), "-"], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=60).stdout
            if name == XLSX_NAME:
                result["xlsx_pages"] = pages
                for token in ("SBER.ME", "Total_Portfolio_Value", "2026-05-26", "Market_Cap"):
                    if token not in extracted:
                        result["errors"].append(f"xlsx_render_missing:{token}")
            else:
                result["docx_pages"] = pages
                for token in ("Portfolio Analysis Report", "Performance Analysis", "Policy Review", "Recommendations"):
                    if token not in extracted:
                        result["errors"].append(f"docx_render_missing:{token}")
            if re.search(r"#(?:NAME|VALUE|REF|DIV/0|N/A)", extracted, flags=re.IGNORECASE):
                result["errors"].append(f"{name}:formula_error_visible")
    if not 5 <= result["xlsx_pages"] <= 8:
        result["errors"].append(f"xlsx_pages={result['xlsx_pages']}")
    if not 2 <= result["docx_pages"] <= 4:
        result["errors"].append(f"docx_pages={result['docx_pages']}")
    return result


def suspicious_content(artifact_dir: Path) -> list[str]:
    findings: list[str] = []
    for path in artifact_dir.iterdir() if artifact_dir.is_dir() else []:
        if not path.is_file():
            continue
        try:
            if path.suffix.lower() in {".md", ".txt", ".eml"}:
                text = path.read_text(encoding="utf-8", errors="ignore").lower()
            elif path.suffix.lower() in {".xlsx", ".docx"}:
                with zipfile.ZipFile(path) as package:
                    text = " ".join(
                        package.read(name).decode("utf-8", errors="ignore")
                        for name in package.namelist() if name.endswith(".xml")
                    ).lower()
            else:
                text = ""
        except Exception:
            continue
        for marker in SUSPICIOUS_MARKERS:
            if marker in text:
                findings.append(f"{path.name}:{marker}")
    return sorted(set(findings))


def criterion(identifier: str, score: float, evidence: object) -> dict[str, object]:
    description, weight, kind, cap = SPECS[identifier]
    raw: dict[str, object] = {"description": description, "weight": weight, "type": kind, "passed": score >= 0.999}
    if cap is not None:
        raw["score_cap_on_failure"] = cap
    return {
        "criterion_id": identifier, "score": round(float(score), 6), "evidence": evidence,
        "method": "deterministic", "raw": raw,
    }


def evaluate(artifact_dir: Path) -> dict[str, object]:
    truth, source_rows, source_errors, summary = load_truth()
    delivered = sorted(path.name for path in artifact_dir.iterdir() if path.is_file()) if artifact_dir.is_dir() else []
    exact_files = delivered == sorted(REQUIRED_FILES)
    packages = {name: inspect_package(artifact_dir / name) for name in (XLSX_NAME, DOCX_NAME)}
    package_ok = all(
        not info["broken"] and not info["missing_targets"] and not info["macros"] and not info["external"]
        for info in packages.values()
    )
    workbook = parse_workbook(artifact_dir / XLSX_NAME, truth, source_rows, summary) if (artifact_dir / XLSX_NAME).is_file() and not source_errors else {
        "readable": False, "identity_errors": ["unavailable"], "boundary_errors": [], "calculation_errors": [],
        "policy_errors": [], "summary_errors": [], "source_errors": [], "formula_errors": [], "native": {},
    }
    word = parse_word(artifact_dir / DOCX_NAME, truth, summary) if (artifact_dir / DOCX_NAME).is_file() and not source_errors else {
        "readable": False, "holding_errors": ["unavailable"], "claim_errors": [], "recommendation_errors": [],
        "unsupported": [], "tables": 0, "headings": [], "recommendation_count": 0,
    }
    email = parse_email(artifact_dir / EML_NAME, artifact_dir) if (artifact_dir / EML_NAME).is_file() else {
        "readable": False, "header_errors": ["unavailable"], "body_errors": [], "attachment_errors": [],
        "unsupported": [], "attachments": [], "body": "",
    }
    render = render_office(artifact_dir) if exact_files else {"xlsx_pages": 0, "docx_pages": 0, "errors": ["incomplete"]}
    suspicious = suspicious_content(artifact_dir)

    identity_ok = workbook.get("readable") and not workbook.get("identity_errors")
    boundary_ok = identity_ok and not workbook.get("boundary_errors")
    calculations_ok = boundary_ok and not workbook.get("calculation_errors") and not workbook.get("formula_errors")
    policy_ok = calculations_ok and not workbook.get("policy_errors")
    summary_ok = calculations_ok and not workbook.get("summary_errors")
    source_trace_ok = workbook.get("readable") and not workbook.get("source_errors")
    word_holdings_ok = word.get("readable") and word.get("tables", 0) >= 3 and not word.get("holding_errors")
    word_claims_ok = word_holdings_ok and not word.get("claim_errors")
    recommendations_ok = word_claims_ok and not word.get("recommendation_errors") and not word.get("unsupported")
    mail_body_ok = email.get("readable") and not email.get("header_errors") and not email.get("body_errors")
    attachments_ok = mail_body_ok and not email.get("attachment_errors")
    native = workbook.get("native", {})
    native_ok = (
        native.get("tables", 0) >= 3 and native.get("charts", 0) >= 1 and native.get("freeze") is True
        and native.get("filters") is True and native.get("conditional_formats", 0) >= 3
        and word.get("tables", 0) >= 3 and not packages[DOCX_NAME].get("media")
    )
    unsupported = sorted(set([*word.get("unsupported", []), *email.get("unsupported", [])]))
    cross_ok = calculations_ok and summary_ok and word_claims_ok and mail_body_ok
    safe_ok = exact_files and package_ok and not suspicious
    render_ok = not render.get("errors")

    results = [
        criterion("R001", 1.0 if exact_files and not source_errors and package_ok else 0.0, {"source_errors": source_errors, "delivered": delivered, "packages": packages}),
        criterion("R002", 1.0 if identity_ok else 0.0, workbook.get("identity_errors")),
        criterion("R003", 1.0 if boundary_ok else 0.0, workbook.get("boundary_errors")),
        criterion("R004", 1.0 if calculations_ok else 0.0, {"calculation_errors": workbook.get("calculation_errors"), "formula_errors": workbook.get("formula_errors")}),
        criterion("R005", 1.0 if policy_ok and not unsupported else 0.0, {"policy_errors": workbook.get("policy_errors"), "unsupported": unsupported}),
        criterion("R006", 1.0 if summary_ok else 0.0, workbook.get("summary_errors")),
        criterion("R007", 1.0 if source_trace_ok else 0.0, workbook.get("source_errors")),
        criterion("R008", 1.0 if word_holdings_ok else 0.0, {"tables": word.get("tables"), "holding_errors": word.get("holding_errors"), "headings": word.get("headings")}),
        criterion("R009", 1.0 if word_claims_ok else 0.0, word.get("claim_errors")),
        criterion("R010", 1.0 if recommendations_ok else 0.0, {"count": word.get("recommendation_count"), "errors": word.get("recommendation_errors"), "unsupported": word.get("unsupported")}),
        criterion("R011", 1.0 if mail_body_ok else 0.0, {"header_errors": email.get("header_errors"), "body_errors": email.get("body_errors")}),
        criterion("R012", 1.0 if attachments_ok else 0.0, {"attachment_errors": email.get("attachment_errors"), "attachments": email.get("attachments")}),
        criterion("R013", 1.0 if native_ok and render_ok else 0.0, {"native": native, "word_tables": word.get("tables"), "render": render}),
        criterion("R014", 1.0 if identity_ok and boundary_ok and calculations_ok and summary_ok and source_trace_ok else 0.0, {"identity": identity_ok, "boundary": boundary_ok, "calculations": calculations_ok, "summary": summary_ok, "source": source_trace_ok}),
        criterion("R015", 1.0 if cross_ok else 0.0, {"workbook": calculations_ok and summary_ok, "word": word_claims_ok, "email": mail_body_ok}),
        criterion("R016", 1.0 if mail_body_ok and attachments_ok else 0.0, {"header_errors": email.get("header_errors"), "attachment_errors": email.get("attachment_errors")}),
        criterion("R017", 1.0 if safe_ok and calculations_ok and native_ok else 0.0, {"safe": safe_ok, "package_ok": package_ok, "suspicious": suspicious, "native": native_ok}),
        criterion("R018", 1.0 if not unsupported else 0.0, unsupported),
    ]

    positive_points = sum(SPECS[item["criterion_id"]][1] * float(item["score"]) for item in results if SPECS[item["criterion_id"]][1] > 0)
    penalty_points = sum(SPECS[item["criterion_id"]][1] for item in results if SPECS[item["criterion_id"]][1] < 0 and float(item["score"]) < 0.999)
    reward = max(0.0, min(1.0, (positive_points + penalty_points) / 100.0))
    hurdle_failures: list[str] = []
    for item in results:
        identifier = item["criterion_id"]
        _, _, kind, cap = SPECS[identifier]
        if kind == "hurdle" and float(item["score"]) < 0.999:
            hurdle_failures.append(identifier)
            reward = min(reward, float(cap))
    reward = round(reward, 6)
    return {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": [sha256(artifact_dir / name) for name in REQUIRED_FILES if (artifact_dir / name).is_file()],
        "criterion_results": results, "reward": reward, "passed": reward >= 0.8 and not hurdle_failures,
        "hurdle_failures": hurdle_failures,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    result = evaluate(args.artifact_dir)
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{result['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
