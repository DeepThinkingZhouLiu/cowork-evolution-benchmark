#!/usr/bin/env python3
"""MOEX 投资组合分析包的任务专属 v2 致命与分层局部反例。"""

from __future__ import annotations

import importlib.util
import json
import shutil
import sys
import tempfile
from email import encoders, policy
from email.parser import BytesParser
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
REQUIRED = ("Portfolio_Analysis.xlsx", "Investment_Memo.docx", "Monthly_Portfolio_Analysis.eml")
XLSX, DOCX, EML = REQUIRED


def _load_judge():
    path = Path(__file__).with_name("judge.py")
    spec = importlib.util.spec_from_file_location("moex_portfolio_v2_negative_judge", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


JUDGE = _load_judge()


def copy_golden(target: Path) -> None:
    target.mkdir(parents=True, exist_ok=True)
    for name in REQUIRED:
        shutil.copy2(GOLDEN / name, target / name)


def sync_email_attachments(root: Path) -> None:
    path = root / EML
    message = BytesParser(policy=policy.default.clone(linesep="\r\n")).parsebytes(path.read_bytes())
    for part in message.iter_attachments():
        name = part.get_filename()
        if name not in (XLSX, DOCX):
            continue
        if part.get("Content-Transfer-Encoding"):
            del part["Content-Transfer-Encoding"]
        part.set_payload((root / name).read_bytes())
        encoders.encode_base64(part)
    path.write_bytes(message.as_bytes())


def mutate_xlsx(root: Path, sheet: str, cell: str, value) -> None:
    path = root / XLSX
    workbook = load_workbook(path, data_only=False)
    workbook[sheet][cell] = value
    workbook.save(path)
    workbook.close()
    sync_email_attachments(root)


def change_doc_paragraph(root: Path, old: str, new: str) -> None:
    path = root / DOCX
    document = Document(path)
    for paragraph in document.paragraphs:
        if old in paragraph.text:
            paragraph.text = paragraph.text.replace(old, new, 1)
            document.save(path)
            sync_email_attachments(root)
            return
    raise AssertionError(f"Word 中未找到：{old}")


def change_doc_table_cell(root: Path, old: str, new: str) -> None:
    path = root / DOCX
    document = Document(path)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if old in cell.text:
                    cell.text = cell.text.replace(old, new, 1)
                    document.save(path)
                    sync_email_attachments(root)
                    return
    raise AssertionError(f"Word 表格中未找到：{old}")


def mutate_mail(root: Path, *, to: str | None = None, old: str | None = None, new: str = "") -> None:
    path = root / EML
    message = BytesParser(policy=policy.default.clone(linesep="\r\n")).parsebytes(path.read_bytes())
    if to is not None:
        del message["To"]
        message["To"] = to
    if old is not None:
        body = message.get_body(preferencelist=("plain",))
        text = body.get_content()
        if old not in text:
            raise AssertionError(f"邮件中未找到：{old}")
        body.set_content(text.replace(old, new, 1), subtype="plain", charset="utf-8")
    path.write_bytes(message.as_bytes())


def holding_company_wrong(root: Path) -> None:
    mutate_xlsx(root, "Holdings", "B2", "错误发行人")


def baseline_price_wrong(root: Path) -> None:
    mutate_xlsx(root, "Holdings", "F2", 126.0)


def current_price_wrong(root: Path) -> None:
    mutate_xlsx(root, "Holdings", "H3", 200.0)


def current_value_formula_wrong(root: Path) -> None:
    mutate_xlsx(root, "Holdings", "I2", "=D2*F2")


def return_formula_wrong(root: Path) -> None:
    mutate_xlsx(root, "Holdings", "J3", "=H3/F2-1")


def policy_formula_flattened(root: Path) -> None:
    mutate_xlsx(root, "Holdings", "P4", "Compliant")


def summary_formula_wrong(root: Path) -> None:
    mutate_xlsx(root, "Summary", "B6", "=SUMIF(Holdings!C2:C6,A6,Holdings!I2:I5)")


def all_formulas_flattened(root: Path) -> None:
    path = root / XLSX
    formulas = load_workbook(path, data_only=False)
    values = load_workbook(path, data_only=True)
    for formula_sheet, value_sheet in zip(formulas.worksheets, values.worksheets, strict=True):
        for row in formula_sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    cell.value = value_sheet[cell.coordinate].value
    formulas.save(path)
    formulas.close()
    values.close()
    sync_email_attachments(root)


def source_price_wrong(root: Path) -> None:
    mutate_xlsx(root, "Source_Prices", "H2", 128.728)


def source_row_missing(root: Path) -> None:
    path = root / XLSX
    workbook = load_workbook(path, data_only=False)
    workbook["Source_Prices"].delete_rows(116)
    workbook.save(path)
    workbook.close()
    sync_email_attachments(root)


def word_holding_price_wrong(root: Path) -> None:
    change_doc_table_cell(root, "120.8171", "121.8171")


def word_best_conflict(root: Path) -> None:
    change_doc_paragraph(root, "MTSS.ME is the best", "GAZP.ME is the best")


def recommendation_loses_evidence(root: Path) -> None:
    change_doc_paragraph(
        root,
        "Monitor SBER.ME before the next committee meeting because it has the weakest observed 30-day return (-5.41%); review new evidence rather than inferring a cause from price movement alone.",
        "Continue routine monitoring.",
    )


def email_wrong_recipient(root: Path) -> None:
    mutate_mail(root, to="wrong-recipient@fund.example.com")


def email_total_conflict(root: Path) -> None:
    mutate_mail(root, old="192,071.02 RUB", new="190,000.00 RUB")


def stale_email_attachment(root: Path) -> None:
    path = root / DOCX
    document = Document(path)
    document.add_paragraph("Committee copy refreshed after packaging.")
    document.save(path)


def upstream_us_data_leak(root: Path) -> None:
    path = root / DOCX
    document = Document(path)
    document.add_paragraph("Legacy upstream row: GOOGL, AMZN, JPM, JNJ, XOM; currency USD.")
    document.save(path)
    sync_email_attachments(root)


def invented_policy_limit(root: Path) -> None:
    path = root / DOCX
    document = Document(path)
    document.add_paragraph("The policy sector limit is 30%, therefore Energy is non-compliant.")
    document.save(path)
    sync_email_attachments(root)


def extra_debug_file(root: Path) -> None:
    (root / "debug.txt").write_text("temporary\n", encoding="utf-8")


PARTIAL = {
    "one_holding_company_wrong": holding_company_wrong,
    "one_baseline_price_wrong": baseline_price_wrong,
    "one_current_price_wrong": current_price_wrong,
    "one_current_value_formula_wrong": current_value_formula_wrong,
    "one_return_formula_wrong_reference": return_formula_wrong,
    "one_policy_formula_flattened": policy_formula_flattened,
    "one_summary_sector_formula_wrong": summary_formula_wrong,
    "all_holdings_and_summary_formulas_flattened": all_formulas_flattened,
    "one_source_price_wrong": source_price_wrong,
    "one_source_row_missing": source_row_missing,
    "word_one_holding_price_wrong": word_holding_price_wrong,
    "word_best_performer_conflict": word_best_conflict,
    "one_recommendation_loses_evidence": recommendation_loses_evidence,
    "email_wrong_recipient": email_wrong_recipient,
    "email_total_value_conflict": email_total_conflict,
    "email_stale_docx_attachment": stale_email_attachment,
    "upstream_us_ticker_and_currency_leak": upstream_us_data_leak,
    "invented_policy_sector_limit": invented_policy_limit,
    "unexpected_debug_file": extra_debug_file,
}


def main() -> None:
    golden = JUDGE.evaluate(GOLDEN)
    assert golden["reward"] == 1.0, golden
    JUDGE._calculated_findings = lambda _path, _truth, _summary: {
        "calculations": [(True, "负例回归复用已验证 Golden 重算") for _ in range(15)],
        "policy": [(True, "负例回归复用已验证 Golden 重算") for _ in range(10)],
        "summary": [(True, "负例回归复用已验证 Golden 重算") for _ in range(10)],
    }
    JUDGE._render_findings = lambda _artifact: [
        (True, "负例回归复用已验证 Golden 渲染"),
        (True, "负例回归复用已验证 Golden Excel 页数"),
        (True, "负例回归复用已验证 Golden Word 页数"),
    ]
    records = []
    with tempfile.TemporaryDirectory(prefix="moex-portfolio-v2-neg-") as raw:
        base = Path(raw)
        for name in REQUIRED:
            for mode in ("missing", "empty", "corrupt"):
                case = base / f"fatal-{mode}-{name}"
                copy_golden(case)
                target = case / name
                if mode == "missing":
                    target.unlink()
                elif mode == "empty":
                    target.write_bytes(b"")
                else:
                    target.write_bytes(b"\xff\xfe\x00")
                reward = JUDGE.evaluate(case)["reward"]
                assert reward == 0.0, (mode, name, reward)
                records.append({"case": f"{mode}:{name}", "class": "fatal", "reward": reward})
        case = base / "fatal-empty-output"
        case.mkdir()
        reward = JUDGE.evaluate(case)["reward"]
        assert reward == 0.0
        records.append({"case": "empty-output", "class": "fatal", "reward": reward})
        case = base / "fatal-input-only"
        shutil.copytree(ROOT / "data" / "input_files", case)
        reward = JUDGE.evaluate(case)["reward"]
        assert reward == 0.0
        records.append({"case": "input-only", "class": "fatal", "reward": reward})

        for name, mutate in PARTIAL.items():
            case = base / f"partial-{name}"
            copy_golden(case)
            mutate(case)
            result = JUDGE.evaluate(case)
            reward = float(result["reward"])
            assert 0.0 < reward < 1.0, (name, reward, result)
            records.append({"case": name, "class": "partial", "reward": reward})

    fatal = [record["reward"] for record in records if record["class"] == "fatal"]
    partial = [record["reward"] for record in records if record["class"] == "partial"]
    print(json.dumps({
        "fatal_count": len(fatal), "fatal_unique_rewards": sorted(set(fatal)),
        "partial_count": len(partial), "partial_min": min(partial), "partial_max": max(partial),
        "cases": records,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
