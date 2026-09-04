#!/usr/bin/env python3
"""逐项攻击本题 Judge，确认其能拒绝行情、公式、政策、邮件与结构错误。"""

from __future__ import annotations

import argparse
import shutil
import tempfile
from email import policy
from email.parser import BytesParser
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from judge import DOCX_NAME, EML_NAME, REQUIRED_FILES, XLSX_NAME, evaluate


CASES = (
    "missing-email",
    "formulas-flattened",
    "wrong-baseline",
    "wrong-source-price",
    "wrong-summary-total",
    "missing-source-row",
    "word-wrong-best",
    "stale-email-attachment",
    "email-wrong-recipient",
    "cross-file-total-conflict",
    "us-ticker-leak",
    "invented-policy-limit",
    "corrupt-docx",
    "extra-debug-file",
)


def copy_golden(golden: Path, case_dir: Path) -> None:
    case_dir.mkdir(parents=True, exist_ok=True)
    for name in REQUIRED_FILES:
        shutil.copy2(golden / name, case_dir / name)


def replace_docx_text(path: Path, old: str, new: str) -> None:
    document = Document(path)
    containers = [document.paragraphs]
    containers.extend(cell.paragraphs for table in document.tables for row in table.rows for cell in row.cells)
    for paragraphs in containers:
        for paragraph in paragraphs:
            if old in paragraph.text:
                paragraph.text = paragraph.text.replace(old, new)
    document.save(path)


def mutate_email(path: Path, *, to: str | None = None, old: str | None = None, new: str | None = None) -> None:
    message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
    if to is not None:
        del message["To"]
        message["To"] = to
    if old is not None:
        body_part = message.get_body(preferencelist=("plain",)) if message.is_multipart() else message
        body_part.set_content(body_part.get_content().replace(old, new or ""), subtype="plain", charset="utf-8")
    path.write_bytes(message.as_bytes(policy=policy.default.clone(linesep="\r\n")))


def mutate(case: str, case_dir: Path) -> None:
    if case == "missing-email":
        (case_dir / EML_NAME).unlink()
    elif case == "formulas-flattened":
        path = case_dir / XLSX_NAME
        formulas = load_workbook(path, data_only=False)
        values = load_workbook(path, data_only=True)
        for formula_sheet, value_sheet in zip(formulas.worksheets, values.worksheets, strict=True):
            for row in formula_sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        cell.value = value_sheet[cell.coordinate].value
        formulas.save(path)
    elif case == "wrong-baseline":
        path = case_dir / XLSX_NAME
        workbook = load_workbook(path, data_only=False)
        workbook["Holdings"]["F2"] = 126.0
        workbook.save(path)
    elif case == "wrong-source-price":
        path = case_dir / XLSX_NAME
        workbook = load_workbook(path, data_only=False)
        workbook["Source_Prices"]["H2"] = float(workbook["Source_Prices"]["H2"].value) + 1.0
        workbook.save(path)
    elif case == "wrong-summary-total":
        path = case_dir / XLSX_NAME
        workbook = load_workbook(path, data_only=False)
        workbook["Summary"]["B2"] = "=SUM(Holdings!I2:I5)"
        workbook.save(path)
    elif case == "missing-source-row":
        path = case_dir / XLSX_NAME
        workbook = load_workbook(path, data_only=False)
        workbook["Source_Prices"].delete_rows(116)
        workbook.save(path)
    elif case == "word-wrong-best":
        replace_docx_text(case_dir / DOCX_NAME, "MTSS.ME", "GAZP.ME")
    elif case == "stale-email-attachment":
        document = Document(case_dir / DOCX_NAME)
        document.add_paragraph("Committee copy refreshed after the email package was created.")
        document.save(case_dir / DOCX_NAME)
    elif case == "email-wrong-recipient":
        mutate_email(case_dir / EML_NAME, to="wrong-recipient@fund.example.com")
    elif case == "cross-file-total-conflict":
        mutate_email(case_dir / EML_NAME, old="192,071.02 RUB", new="190,000.00 RUB")
    elif case == "us-ticker-leak":
        document = Document(case_dir / DOCX_NAME)
        document.add_paragraph("Legacy upstream row: GOOGL, AMZN, JPM, JNJ, XOM; currency USD.")
        document.save(case_dir / DOCX_NAME)
    elif case == "invented-policy-limit":
        document = Document(case_dir / DOCX_NAME)
        document.add_paragraph("The policy sector limit is 30%, therefore Energy is non-compliant.")
        document.save(case_dir / DOCX_NAME)
    elif case == "corrupt-docx":
        (case_dir / DOCX_NAME).write_bytes(b"not-a-docx")
    elif case == "extra-debug-file":
        (case_dir / "debug.txt").write_text("temporary judge notes", encoding="utf-8")
    else:
        raise ValueError(case)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--case", choices=CASES, action="append", required=True)
    args = parser.parse_args()
    failures: list[dict[str, object]] = []
    for case in args.case:
        with tempfile.TemporaryDirectory(prefix=f"moex-portfolio-negative-{case}-") as tmp:
            case_dir = Path(tmp)
            copy_golden(args.golden, case_dir)
            mutate(case, case_dir)
            result = evaluate(case_dir)
            print(f"{case}\treward={result['reward']:.6f}\tpassed={result['passed']}\thurdles={','.join(result['hurdle_failures'])}")
            if result["passed"] or float(result["reward"]) >= 0.8:
                failures.append({"case": case, "result": result})
    if failures:
        raise SystemExit(f"负例错误通过：{[item['case'] for item in failures]}")


if __name__ == "__main__":
    main()
