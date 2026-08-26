#!/usr/bin/env python3
"""生成本题的六个 Golden 交付物。"""

from __future__ import annotations

import argparse
import json
import shutil
from math import ceil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from refund_root_cause_processor import analyze


NAVY = "0B1F33"
TEAL = "167D95"
TEAL_LIGHT = "DCEFF2"
ICE = "F4F7FA"
ORANGE = "D97706"
RED = "B42318"
GRAY = "5B6673"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D7DEE5", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def add_table(doc: Document, headers: list[str], rows: list[list[object]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, NAVY)
        set_cell_border(cell, NAVY)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in paragraph.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(8.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
        if widths:
            cell.width = Inches(widths[index])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.text = str(value)
            set_cell_shading(cell, WHITE if row_index % 2 == 0 else ICE)
            set_cell_border(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(8)
                    run.font.color.rgb = RGBColor.from_string(NAVY)
            if widths:
                cell.width = Inches(widths[index])
    return table


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(NAVY)


def build_excel(result: dict, output_path: Path) -> None:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)

    header_fill = PatternFill("solid", fgColor=NAVY)
    sub_fill = PatternFill("solid", fgColor=TEAL_LIGHT)
    header_font = Font(name="Aptos", bold=True, color=WHITE, size=10)
    body_font = Font(name="Aptos", color=NAVY, size=10)
    border = Border(bottom=Side(style="thin", color="D7DEE5"))

    detail_headers = [
        "Refund_ID",
        "Order_ID",
        "Product_Name",
        "Refund_Amount",
        "Severity",
        "Refund_Reason",
        "Date",
    ]
    detail_ws = workbook.create_sheet("Refund Details")
    detail_ws.append(detail_headers)
    for row in result["refund_details"]:
        detail_ws.append([row[key] for key in detail_headers])
    detail_ws.freeze_panes = "A2"
    detail_ws.auto_filter.ref = detail_ws.dimensions
    detail_table = Table(displayName="RefundDetailsTable", ref=detail_ws.dimensions)
    detail_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False
    )
    detail_ws.add_table(detail_table)
    for cell in detail_ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row_number, row in enumerate(detail_ws.iter_rows(min_row=2), start=2):
        for cell in row:
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[3].number_format = '$#,##0.00'
        product_text = str(row[2].value or "")
        detail_ws.row_dimensions[row_number].height = max(24, min(120, 15 * ceil(len(product_text) / 68)))
    detail_ws.column_dimensions["A"].width = 12
    detail_ws.column_dimensions["B"].width = 12
    detail_ws.column_dimensions["C"].width = 70
    detail_ws.column_dimensions["D"].width = 18
    detail_ws.column_dimensions["E"].width = 14
    detail_ws.column_dimensions["F"].width = 25
    detail_ws.column_dimensions["G"].width = 14
    detail_ws.row_dimensions[1].height = 30
    detail_ws.sheet_view.showGridLines = False
    detail_ws.page_setup.orientation = "landscape"
    detail_ws.page_setup.paperSize = detail_ws.PAPERSIZE_A4
    detail_ws.page_setup.fitToWidth = 1
    detail_ws.page_setup.fitToHeight = 0
    detail_ws.sheet_properties.pageSetUpPr.fitToPage = True
    detail_ws.print_title_rows = "1:1"
    detail_ws.print_area = f"A1:G{detail_ws.max_row}"
    detail_ws.page_margins.left = 0.25
    detail_ws.page_margins.right = 0.25
    detail_ws.page_margins.top = 0.45
    detail_ws.page_margins.bottom = 0.45
    detail_ws.conditional_formatting.add(
        f"D2:D{detail_ws.max_row}",
        CellIsRule(operator="greaterThan", formula=["50"], fill=PatternFill("solid", fgColor="FDE7E5")),
    )

    impact_headers = [
        "Product_Name",
        "Refund_Count",
        "Total_Refund_Amount",
        "Avg_Review_Rating",
        "Review_Count",
        "Issue_Type",
        "Requires_Investigation",
    ]
    impact_ws = workbook.create_sheet("Product Impact")
    impact_ws.append(impact_headers)
    for row in result["product_impact"]:
        impact_ws.append([row[key] for key in impact_headers])
    impact_ws.freeze_panes = "A2"
    impact_ws.auto_filter.ref = impact_ws.dimensions
    impact_table = Table(displayName="ProductImpactTable", ref=impact_ws.dimensions)
    impact_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False
    )
    impact_ws.add_table(impact_table)
    for cell in impact_ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row_number, row in enumerate(impact_ws.iter_rows(min_row=2), start=2):
        for cell in row:
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[2].number_format = '$#,##0.00'
        row[3].number_format = '0.00'
        product_text = str(row[0].value or "")
        impact_ws.row_dimensions[row_number].height = max(23, min(64, 15 * ceil(len(product_text) / 68)))
    widths = {"A": 72, "B": 15, "C": 23, "D": 22, "E": 16, "F": 20, "G": 24}
    for column, width in widths.items():
        impact_ws.column_dimensions[column].width = width
    impact_ws.row_dimensions[1].height = 34
    impact_ws.sheet_view.showGridLines = False
    impact_ws.page_setup.orientation = "landscape"
    impact_ws.page_setup.paperSize = impact_ws.PAPERSIZE_A4
    impact_ws.page_setup.fitToWidth = 1
    impact_ws.page_setup.fitToHeight = 0
    impact_ws.sheet_properties.pageSetUpPr.fitToPage = True
    impact_ws.print_title_rows = "1:1"
    impact_ws.print_area = f"A1:G{impact_ws.max_row}"
    impact_ws.page_margins.left = 0.25
    impact_ws.page_margins.right = 0.25
    impact_ws.page_margins.top = 0.45
    impact_ws.page_margins.bottom = 0.45
    impact_ws.conditional_formatting.add(
        f"B2:B{impact_ws.max_row}",
        CellIsRule(operator="greaterThanOrEqual", formula=["2"], fill=PatternFill("solid", fgColor="FFF1D6")),
    )

    summary_ws = workbook.create_sheet("Summary")
    summary_ws.append(["Metric", "Value"])
    for key, value in result["summary"].items():
        summary_ws.append([key, value])
    summary_ws.freeze_panes = "A2"
    summary_table = Table(displayName="SummaryTable", ref=summary_ws.dimensions)
    summary_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False
    )
    summary_ws.add_table(summary_table)
    for cell in summary_ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    for row in summary_ws.iter_rows(min_row=2):
        row[0].font = Font(name="Aptos", bold=True, color=NAVY)
        row[0].fill = sub_fill
        row[1].font = Font(name="Aptos Display", bold=True, color=TEAL, size=12)
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")
    summary_ws["B3"].number_format = '$#,##0.00'
    summary_ws.column_dimensions["A"].width = 42
    summary_ws.column_dimensions["B"].width = 22
    summary_ws.row_dimensions[1].height = 30
    summary_ws.sheet_view.showGridLines = False
    summary_ws.page_setup.orientation = "portrait"
    summary_ws.page_setup.paperSize = summary_ws.PAPERSIZE_A4
    summary_ws.page_setup.fitToWidth = 1
    summary_ws.page_setup.fitToHeight = 1
    summary_ws.sheet_properties.pageSetUpPr.fitToPage = True
    summary_ws.print_area = f"A1:B{summary_ws.max_row}"

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(output_path)


def build_report(result: dict, output_path: Path) -> None:
    summary = result["summary"]
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].font.color.rgb = RGBColor.from_string(NAVY)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for style_name, size, color in [
        ("Title", 30, NAVY),
        ("Subtitle", 12, TEAL),
        ("Heading 1", 19, NAVY),
        ("Heading 2", 12, TEAL),
    ]:
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.text = "QUALITY CONTROL  /  REFUND INVESTIGATION"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(TEAL)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Source-backed analysis  •  Cowork Evolution Benchmark")
    footer_run.font.name = "Aptos"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string(GRAY)

    title = doc.add_paragraph(style="Title")
    title.add_run("Refund Root Cause\nAnalysis")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Cross-file quality-control package")
    doc.add_paragraph(
        "Prepared from the provided refund, order, product-review and policy inputs. "
        "All monetary values use the source currency and all classifications follow Quality_Policy.pdf."
    )

    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(
        f"The source contains {summary['Total_Refunds']} refunds totaling "
        f"${summary['Total_Refund_Amount']:,.2f} across {summary['Products_Affected']} returned products. "
        f"Severity is {summary['Critical_Count']} Critical, {summary['Major_Count']} Major and "
        f"{summary['Minor_Count']} Minor. {summary['Products_Requiring_Investigation']} products meet "
        "the policy threshold of two or more refunds and require supplier investigation."
    )
    add_table(
        doc,
        ["Measure", "Result", "Measure", "Result"],
        [
            ["Refunds", summary["Total_Refunds"], "Refund amount", f"${summary['Total_Refund_Amount']:,.2f}"],
            ["Products affected", summary["Products_Affected"], "Supplier investigations", summary["Products_Requiring_Investigation"]],
            ["Quality Issues", summary["Quality_Issues"], "Service Issues", summary["Service_Issues"]],
        ],
        [1.45, 1.0, 1.7, 1.0],
    )

    doc.add_page_break()
    doc.add_heading("Refund Trend Analysis", level=1)
    doc.add_paragraph(
        "The breakdown below counts the stated refund reasons exactly as supplied. "
        "The input is a cross-sectional refund snapshot, so no unsupported time-series direction is asserted."
    )
    add_table(
        doc,
        ["Refund reason", "Refund count"],
        [[row["Refund_Reason"], row["Refund_Count"]] for row in result["refund_reason_breakdown"]],
        [3.8, 1.2],
    )
    doc.add_heading("Root Cause Classification", level=1)
    doc.add_paragraph(
        "Reasons are mapped using guide.md: defective products to Manufacturing Defect; damaged arrivals to "
        "Shipping Damage; wrong-item cases to Wrong Item; changed-mind or not-as-described cases to Customer "
        "Expectations; duplicate orders to Other."
    )
    add_table(
        doc,
        ["Root cause", "Refund count"],
        [[row["Root_Cause"], row["Refund_Count"]] for row in result["root_cause_breakdown"]],
        [3.8, 1.2],
    )

    doc.add_page_break()
    doc.add_heading("Product-Level Findings", level=1)
    doc.add_paragraph(
        "The four products below meet the two-refund supplier-investigation threshold. Product refund amounts "
        "allocate each order-level refund equally across that order's returned line items, then round the "
        "product total to two decimals."
    )
    add_table(
        doc,
        ["Product", "Refunds", "Allocated amount", "Avg rating", "Issue"],
        [
            [
                row["Product_Name"],
                row["Refund_Count"],
                f"${row['Total_Refund_Amount']:,.2f}",
                f"{row['Avg_Review_Rating']:.2f}" if row["Avg_Review_Rating"] is not None else "No reviews",
                row["Issue_Type"],
            ]
            for row in result["investigation_products"]
        ],
        [2.8, 0.55, 0.9, 0.7, 0.85],
    )
    doc.add_paragraph(
        f"Across all {summary['Products_Affected']} affected products, the policy classifies "
        f"{summary['Quality_Issues']} as Quality Issue, {summary['Service_Issues']} as Service Issue and "
        f"{summary['Products_Affected'] - summary['Quality_Issues'] - summary['Service_Issues']} as No Reviews."
    )

    doc.add_page_break()
    doc.add_heading("Recommendations", level=1)
    add_bullet(doc, "Open supplier investigations for the four threshold products and attach the Product Impact evidence from Refund_Analysis.xlsx.")
    add_bullet(doc, "Review duplicate-order controls because Duplicate order is the largest stated reason, with four refunds.")
    add_bullet(doc, "Review packaging and carrier handoff evidence for the two Product arrived damaged refunds.")
    add_bullet(doc, "Review catalog accuracy and fulfillment checks for the two Item not as described and two Wrong item received refunds.")
    add_bullet(doc, "Continue review-rating monitoring; this snapshot contains no policy-defined Quality Issue because every affected product with reviews averages at least 3.0.")

    doc.add_heading("Action Plan", level=1)
    action_rows = [
        ["Immediate", "Quality team", "Validate all 10 Critical refunds and open the four supplier cases.", "Case list matches Product Impact."],
        ["Short term", "Fulfillment", "Review duplicate-order, wrong-item and damaged-arrival controls.", "Findings cite the supplied reason records."],
        ["Long term", "Quality operations", "Monitor refund counts and average review ratings with the same policy thresholds.", "Recurring analysis remains replayable from source exports."],
    ]
    add_table(doc, ["Horizon", "Owner", "Action", "Evidence of completion"], action_rows, [0.8, 1.0, 2.3, 1.45])
    doc.add_paragraph(
        "Evidence boundary: the supplied files support refund counts, amounts, reasons, product allocations and review-based issue types. "
        "They do not support revenue impact, customer-loss estimates, supplier fault rates, industry benchmarks or trend direction."
    )

    doc.core_properties.title = "Refund Root Cause Analysis"
    doc.core_properties.subject = "Source-backed cross-file refund investigation"
    doc.core_properties.author = "Quality Control"
    doc.save(output_path)


def build_emails(result: dict, output_dir: Path) -> None:
    summary = result["summary"]
    reasons = ", ".join(
        f"{row['Refund_Reason']} {row['Refund_Count']}"
        for row in result["refund_reason_breakdown"]
    )
    quality_text = f"""To: quality_team@company.com
Subject: Refund Analysis — Quality Control Findings

Hello Quality Team,

The provided source snapshot contains {summary['Total_Refunds']} refunds totaling ${summary['Total_Refund_Amount']:,.2f} across {summary['Products_Affected']} affected products.

Severity breakdown: Critical {summary['Critical_Count']}; Major {summary['Major_Count']}; Minor {summary['Minor_Count']}.
Refund reasons: {reasons}.
Policy classification across affected products: Quality Issue {summary['Quality_Issues']}; Service Issue {summary['Service_Issues']}; products requiring supplier investigation {summary['Products_Requiring_Investigation']}.

The detailed evidence is in Refund_Analysis.xlsx and Root_Cause_Report.docx. This is an offline draft; no message has been sent.

Quality Control
"""
    (output_dir / "quality_team_email.txt").write_text(quality_text, encoding="utf-8")

    product_lines = "\n".join(
        f"- {row['Product_Name']} | refunds: {row['Refund_Count']} | allocated refund amount: ${row['Total_Refund_Amount']:,.2f}"
        for row in result["investigation_products"]
    )
    supplier_text = f"""To: supplier_relations@company.com
Subject: Supplier Investigation — Refund Threshold Products

Hello Supplier Relations,

The Quality Policy requires investigation when a product has two or more refunds. The following {summary['Products_Requiring_Investigation']} products meet that threshold:

{product_lines}

Counts and amounts match the Product Impact sheet in Refund_Analysis.xlsx. Amounts allocate each order-level refund equally across the order's returned products. This is an offline draft; no message has been sent.

Quality Control
"""
    (output_dir / "supplier_relations_email.txt").write_text(supplier_text, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()
    input_dir = args.input_dir.resolve()
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    result = analyze(input_dir)
    (output_dir / "analysis_results.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    build_excel(result, output_dir / "Refund_Analysis.xlsx")
    build_report(result, output_dir / "Root_Cause_Report.docx")
    build_emails(result, output_dir)
    processor_source = Path(__file__).with_name("refund_root_cause_processor.py")
    processor_target = output_dir / "refund_root_cause_processor.py"
    if processor_source.resolve() != processor_target.resolve():
        shutil.copy2(processor_source, processor_target)


if __name__ == "__main__":
    main()
