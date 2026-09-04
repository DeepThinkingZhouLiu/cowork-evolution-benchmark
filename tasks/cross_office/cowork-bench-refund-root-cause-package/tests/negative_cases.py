#!/usr/bin/env python3
"""退款根因六文件包的任务专属 v2 分层负例回归。"""

from __future__ import annotations

import importlib.util
import json
import os
import shutil
import tempfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


TASK_ROOT = Path(__file__).resolve().parents[1]
GOLDEN = TASK_ROOT / "solution"
CORE = [
    "refund_root_cause_processor.py",
    "analysis_results.json",
    "Refund_Analysis.xlsx",
    "Root_Cause_Report.docx",
    "quality_team_email.txt",
    "supplier_relations_email.txt",
]


def load_judge():
    path = Path(__file__).with_name("judge.py")
    spec = importlib.util.spec_from_file_location("refund_judge_v2_negative", path)
    if spec is None or spec.loader is None:
        raise RuntimeError(path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


JUDGE = load_judge()


def seed(target: Path) -> None:
    target.mkdir(parents=True, exist_ok=True)
    for name in CORE:
        shutil.copy2(GOLDEN / name, target / name)


def mutate_json(target: Path, callback) -> None:
    path = target / "analysis_results.json"
    payload = json.loads(path.read_text(encoding="utf-8"))
    callback(payload)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def mutate_xlsx(target: Path, callback) -> None:
    path = target / "Refund_Analysis.xlsx"
    workbook = load_workbook(path)
    callback(workbook)
    workbook.save(path)
    workbook.close()


def replace_docx(target: Path, old: str, new: str) -> None:
    path = target / "Root_Cause_Report.docx"
    document = Document(path)
    for paragraph in document.paragraphs:
        if old in paragraph.text:
            for run in paragraph.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
                    document.save(path)
                    return
            paragraph.text = paragraph.text.replace(old, new)
            document.save(path)
            return
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if old in cell.text:
                    cell.text = cell.text.replace(old, new)
                    document.save(path)
                    return
    raise AssertionError(f"DOCX 中找不到待替换文本：{old}")


def replace_text(target: Path, name: str, old: str, new: str) -> None:
    path = target / name
    text = path.read_text(encoding="utf-8")
    if old not in text:
        raise AssertionError(f"{name} 中找不到：{old}")
    path.write_text(text.replace(old, new, 1), encoding="utf-8")


def case(name: str, kind: str, mutator) -> dict:
    with tempfile.TemporaryDirectory(prefix=f"refund-negative-{name}-") as temp_dir:
        artifact_dir = Path(temp_dir) / "output"
        seed(artifact_dir)
        mutator(artifact_dir)
        os.environ["COWORK_EVO_SKIP_RENDER_FOR_TESTS"] = "1"
        result = JUDGE.evaluate_v2(artifact_dir)
    reward = result["reward"]
    if kind == "fatal":
        assert reward == 0.0, (name, reward, result.get("hard_zero_reasons"))
    else:
        assert 0.0 < reward < 1.0, (name, reward)
    return {"case": name, "class": kind, "reward": reward}


def empty_output(target: Path) -> None:
    for path in target.iterdir():
        path.unlink()


def missing(name: str):
    return lambda target: (target / name).unlink()


def zero_byte(name: str):
    return lambda target: (target / name).write_bytes(b"")


def corrupt(name: str):
    return lambda target: (target / name).write_bytes(b"not-a-valid-office-file")


def hardcoded_processor(target: Path) -> None:
    payload = (target / "analysis_results.json").read_text(encoding="utf-8")
    script = f'''#!/usr/bin/env python3
import argparse
from pathlib import Path

PAYLOAD = {payload!r}
parser = argparse.ArgumentParser()
parser.add_argument("--input-dir", required=True)
parser.add_argument("--output", required=True)
args = parser.parse_args()
Path(args.output).write_text(PAYLOAD, encoding="utf-8")
'''
    (target / "refund_root_cause_processor.py").write_text(script, encoding="utf-8")


def shell_payload() -> dict:
    return {
        "refund_details": [],
        "product_impact": [],
        "summary": {"Total_Refunds": 1},
        "refund_reason_breakdown": [],
        "root_cause_breakdown": [],
        "investigation_products": [],
    }


def shell_processor(target: Path) -> None:
    payload = json.dumps(shell_payload(), ensure_ascii=False, indent=2) + "\n"
    script = f'''#!/usr/bin/env python3
import argparse
from pathlib import Path

PAYLOAD = {payload!r}
parser = argparse.ArgumentParser()
parser.add_argument("--input-dir", required=True)
parser.add_argument("--output", required=True)
args = parser.parse_args()
Path(args.output).write_text(PAYLOAD, encoding="utf-8")
'''
    (target / "refund_root_cause_processor.py").write_text(script, encoding="utf-8")


def layered_valid_partial(target: Path) -> None:
    shell_processor(target)
    (target / "analysis_results.json").write_text(
        json.dumps(shell_payload(), ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    def mutate(workbook) -> None:
        workbook["Product Impact"].delete_rows(14, 12)
    mutate_xlsx(target, mutate)


def conflicting_summary(target: Path) -> None:
    mutate_json(target, lambda payload: payload["summary"].update(
        Total_Refunds=10, Total_Refund_Amount=629.51,
        Critical_Count=8, Major_Count=3, Minor_Count=2,
    ))
    def mutate(workbook) -> None:
        values = [11, 1221.32, 7, 4, 2]
        for row, value in zip(range(2, 7), values):
            workbook["Summary"].cell(row, 2).value = value
    mutate_xlsx(target, mutate)
    replace_docx(target, "$4,407.83", "$1,680.25")
    replace_text(target, "quality_team_email.txt", "$4,407.83", "$88.03")


def canonical_processor_error(target: Path) -> None:
    replace_text(
        target,
        "refund_root_cause_processor.py",
        'amount = Decimal(refund["amount"])',
        'amount = Decimal(refund["amount"]) + (Decimal("1.00") if refund_id == 4 else Decimal("0"))',
    )


def remove_supplier_line(target: Path) -> None:
    path = target / "supplier_relations_email.txt"
    lines = path.read_text(encoding="utf-8").splitlines()
    needle = "MICROSMT Smart LED TV"
    path.write_text("\n".join(line for line in lines if needle not in line) + "\n", encoding="utf-8")


def reorder_excel_products(target: Path) -> None:
    def mutate(workbook) -> None:
        sheet = workbook["Product Impact"]
        left = [sheet.cell(2, column).value for column in range(1, 8)]
        right = [sheet.cell(3, column).value for column in range(1, 8)]
        for column, value in enumerate(right, 1):
            sheet.cell(2, column).value = value
        for column, value in enumerate(left, 1):
            sheet.cell(3, column).value = value
    mutate_xlsx(target, mutate)


def reorder_json_details(target: Path) -> None:
    mutate_json(target, lambda payload: payload["refund_details"].reverse())


def wrong_word_root_count(target: Path) -> None:
    path = target / "Root_Cause_Report.docx"
    document = Document(path)
    for table in document.tables:
        for row in table.rows:
            if row.cells and row.cells[0].text.strip() == "Other":
                row.cells[1].text = "3"
                document.save(path)
                return
    raise AssertionError("Word 根因表中找不到 Other")


def main() -> None:
    fatal_cases = [
        ("empty-output", empty_output),
        *[(f"missing-{name}", missing(name)) for name in CORE],
        ("zero-byte-json", zero_byte("analysis_results.json")),
        ("corrupt-xlsx", corrupt("Refund_Analysis.xlsx")),
        ("corrupt-docx", corrupt("Root_Cause_Report.docx")),
        ("invalid-processor", lambda target: (target / "refund_root_cause_processor.py").write_text("def broken(:\n", encoding="utf-8")),
        ("empty-quality-email", zero_byte("quality_team_email.txt")),
        ("input-renamed-as-json", lambda target: shutil.copy2(TASK_ROOT / "data/input_files/refunds.csv", target / "analysis_results.json")),
    ]
    partial_cases = [
        ("json-one-refund-amount", lambda target: mutate_json(target, lambda payload: payload["refund_details"][3].update(Refund_Amount=16.45))),
        ("json-one-product-rating", lambda target: mutate_json(target, lambda payload: payload["product_impact"][0].update(Avg_Review_Rating=1.0))),
        ("json-one-summary-value", lambda target: mutate_json(target, lambda payload: payload["summary"].update(Critical_Count=9))),
        ("json-one-product-missing", lambda target: mutate_json(target, lambda payload: payload["product_impact"].pop())),
        ("json-one-root-count", lambda target: mutate_json(target, lambda payload: payload["root_cause_breakdown"][0].update(Refund_Count=3))),
        ("json-details-reversed", reorder_json_details),
        ("excel-one-detail-amount", lambda target: mutate_xlsx(target, lambda workbook: setattr(workbook["Refund Details"]["D5"], "value", 16.45))),
        ("excel-one-severity", lambda target: mutate_xlsx(target, lambda workbook: setattr(workbook["Refund Details"]["E5"], "value", "Minor"))),
        ("excel-one-product-amount", lambda target: mutate_xlsx(target, lambda workbook: setattr(workbook["Product Impact"]["C2"], "value", 331.0))),
        ("excel-one-product-rating", lambda target: mutate_xlsx(target, lambda workbook: setattr(workbook["Product Impact"]["D2"], "value", 1.0))),
        ("excel-product-order-swap", reorder_excel_products),
        ("excel-one-summary-value", lambda target: mutate_xlsx(target, lambda workbook: setattr(workbook["Summary"]["B4"], "value", 9))),
        ("excel-header-typo", lambda target: mutate_xlsx(target, lambda workbook: setattr(workbook["Refund Details"]["D1"], "value", "Refund_Value"))),
        ("word-one-heading-missing", lambda target: replace_docx(target, "Refund Trend Analysis", "Refund Review")),
        ("word-total-wrong", lambda target: replace_docx(target, "$4,407.83", "$4,307.83")),
        ("word-root-count-wrong", wrong_word_root_count),
        ("word-investigation-product-missing", lambda target: replace_docx(target, "MICROSMT Smart LED TV", "MICROSMT display")),
        ("quality-recipient-wrong", lambda target: replace_text(target, "quality_team_email.txt", "quality_team@company.com", "support@company.com")),
        ("quality-severity-wrong", lambda target: replace_text(target, "quality_team_email.txt", "Critical 10", "Critical 9")),
        ("quality-reason-count-wrong", lambda target: replace_text(target, "quality_team_email.txt", "Duplicate order 4", "Duplicate order 3")),
        ("supplier-one-product-missing", remove_supplier_line),
        ("supplier-one-amount-wrong", lambda target: replace_text(target, "supplier_relations_email.txt", "$332.00", "$331.00")),
        ("supplier-subject-wrong", lambda target: replace_text(target, "supplier_relations_email.txt", "Supplier Investigation", "Vendor Review")),
        ("processor-hardcoded", hardcoded_processor),
        ("processor-canonical-error", canonical_processor_error),
        ("multi-component-valid-partial", layered_valid_partial),
        ("cross-artifact-summary-conflict", conflicting_summary),
        ("unsupported-roi", lambda target: (target / "quality_team_email.txt").write_text((target / "quality_team_email.txt").read_text(encoding="utf-8") + "Projected ROI is 25%.\n", encoding="utf-8")),
        ("unsupported-money", lambda target: (target / "quality_team_email.txt").write_text((target / "quality_team_email.txt").read_text(encoding="utf-8") + "Projected revenue loss is $999.00.\n", encoding="utf-8")),
        ("unexpected-debug-file", lambda target: (target / "debug.log").write_text("debug\n", encoding="utf-8")),
    ]

    records = [case(name, "fatal", mutator) for name, mutator in fatal_cases]
    records.extend(case(name, "partial", mutator) for name, mutator in partial_cases)
    fatal_rewards = [record["reward"] for record in records if record["class"] == "fatal"]
    partial_rewards = [record["reward"] for record in records if record["class"] == "partial"]
    print(json.dumps({
        "fatal_count": len(fatal_rewards),
        "fatal_unique_rewards": sorted(set(fatal_rewards)),
        "partial_count": len(partial_rewards),
        "partial_min": min(partial_rewards),
        "partial_max": max(partial_rewards),
        "cases": records,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
