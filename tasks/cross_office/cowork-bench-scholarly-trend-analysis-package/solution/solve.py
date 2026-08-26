#!/usr/bin/env python3
"""Cowork-Bench 学术趋势抽取、分析与呈现题的逐题参考解法。"""

from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


PROCESSOR_NAME = "scholarly_trend_processor.py"
JSON_NAME = "scholarly_trend_results.json"
XLSX_NAME = "Trend_Analysis_Report.xlsx"
PPTX_NAME = "Trend_Analysis_Presentation.pptx"

NAVY = RGBColor(20, 45, 78)
BLUE = RGBColor(45, 118, 182)
CYAN = RGBColor(65, 182, 196)
CORAL = RGBColor(239, 111, 92)
GOLD = RGBColor(245, 190, 65)
INK = RGBColor(31, 41, 55)
MUTED = RGBColor(91, 107, 127)
PALE = RGBColor(235, 243, 250)
WHITE = RGBColor(255, 255, 255)


def add_excel_table(sheet, ref: str, name: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
    sheet.add_table(table)


def style_sheet(sheet, widths: dict[str, float]) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 28
    thin = Side(style="thin", color="D9E2F3")
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=10, color="1F2937")
            cell.alignment = Alignment(vertical="center", wrap_text=True, horizontal="left" if cell.column <= 3 else "right")
            cell.border = Border(bottom=thin)
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.oddFooter.center.text = "Cowork-Bench · Scholarly Trend Analysis"


def write_workbook(path: Path, result: dict) -> None:
    workbook = Workbook()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.properties.creator = "Cowork Evolution Benchmark"
    workbook.properties.title = "Scholarly Trend Analysis"

    data_sheet = workbook.active
    data_sheet.title = "Data_Analysis"
    data_headers = ["Paper_ID", "Title", "Area", "Citations", "Relevance_Score"]
    data_sheet.append(data_headers)
    for item in result["data_analysis"]:
        data_sheet.append([item[header] for header in data_headers])
    style_sheet(data_sheet, {"A": 18, "B": 36, "C": 30, "D": 16, "E": 20})
    for row in range(2, 6):
        data_sheet.cell(row, 4).number_format = "0.0"
        data_sheet.cell(row, 5).number_format = "0"
    add_excel_table(data_sheet, "A1:E5", "Scholarly_Data_Analysis")
    data_sheet.conditional_formatting.add(
        "E2:E5",
        ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"),
    )

    metrics_sheet = workbook.create_sheet("Metrics")
    metrics_sheet.append(["Metric", "Value"])
    metrics_sheet.append(["Total_Papers", "=COUNTA(Data_Analysis!A2:A5)"])
    metrics_sheet.append(["Avg_Citations", "=ROUND(AVERAGE(Data_Analysis!D2:D5),0)"])
    metrics_sheet.append(["Top_Area", "=INDEX(Data_Analysis!C2:C5,MATCH(MAX(Data_Analysis!E2:E5),Data_Analysis!E2:E5,0))"])
    metrics_sheet.append(["Avg_Relevance", "=ROUND(AVERAGE(Data_Analysis!E2:E5),0)"])
    style_sheet(metrics_sheet, {"A": 24, "B": 34})
    metrics_sheet["B2"].number_format = "0"
    metrics_sheet["B3"].number_format = "0"
    metrics_sheet["B5"].number_format = "0"
    add_excel_table(metrics_sheet, "A1:B5", "Scholarly_Metrics")

    recommendation_sheet = workbook.create_sheet("Recommendations")
    recommendation_headers = ["Priority", "Action", "Area"]
    recommendation_sheet.append(recommendation_headers)
    for item in result["recommendations"]:
        recommendation_sheet.append([item[header] for header in recommendation_headers])
    style_sheet(recommendation_sheet, {"A": 16, "B": 68, "C": 30})
    add_excel_table(recommendation_sheet, f"A1:C{recommendation_sheet.max_row}", "Scholarly_Recommendations")
    workbook.save(path)


def add_text(slide, left, top, width, height, text, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT, font="Arial"):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Pt(0)
    frame.margin_right = Pt(0)
    frame.margin_top = Pt(0)
    frame.margin_bottom = Pt(0)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_rect(slide, left, top, width, height, fill, radius=True, line=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
    return shape


def add_title(slide, title: str, kicker: str, number: int) -> None:
    add_text(slide, 0.65, 0.28, 1.8, 0.26, kicker.upper(), 9, BLUE, True)
    add_text(slide, 0.65, 0.62, 11.6, 0.58, title, 26, NAVY, True)
    add_text(slide, 12.25, 0.35, 0.42, 0.3, f"{number:02d}", 10, MUTED, True, PP_ALIGN.RIGHT)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(1.25), Inches(12.0), Inches(0.025))
    line.fill.solid(); line.fill.fore_color.rgb = RGBColor(213, 224, 235); line.line.fill.background()


def add_footer(slide, text: str = "Источники: scholarly_papers.json + dashboard.html · срез март 2026") -> None:
    add_text(slide, 0.65, 7.12, 12.0, 0.22, text, 8, MUTED)


def add_card(slide, left, top, width, height, title, body, accent=BLUE) -> None:
    add_rect(slide, left, top, width, height, WHITE, True, RGBColor(213, 224, 235))
    add_rect(slide, left, top, 0.08, height, accent, False)
    add_text(slide, left + 0.25, top + 0.18, width - 0.45, 0.35, title, 14, NAVY, True)
    add_text(slide, left + 0.25, top + 0.65, width - 0.45, height - 0.82, body, 10, INK)


def write_presentation(path: Path, result: dict) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    slide = presentation.slides.add_slide(blank)
    background = slide.background.fill
    background.solid(); background.fore_color.rgb = RGBColor(246, 249, 252)
    add_rect(slide, 0, 0, 5.0, 7.5, NAVY, False)
    add_text(slide, 0.7, 0.7, 3.8, 0.35, "SCHOLARLY TREND", 11, CYAN, True)
    add_text(slide, 0.7, 1.35, 3.8, 1.45, "Научные\nтренды", 34, WHITE, True)
    add_text(slide, 0.7, 3.15, 3.7, 0.9, "Сопоставление статей\nс внешними метриками", 16, WHITE)
    add_text(slide, 0.7, 6.55, 3.7, 0.4, "Март 2026 · исследовательская группа", 9, RGBColor(197, 214, 232))
    for left, value, label, accent in ((5.65, "4", "статьи", BLUE), (8.0, "4", "области", CYAN), (10.35, "95", "макс. score", CORAL)):
        add_rect(slide, left, 1.15, 1.8, 1.55, WHITE, True, RGBColor(213, 224, 235))
        add_text(slide, left + 0.2, 1.35, 1.4, 0.6, value, 29, accent, True, PP_ALIGN.CENTER)
        add_text(slide, left + 0.15, 2.08, 1.5, 0.3, label, 10, MUTED, False, PP_ALIGN.CENTER)
    add_text(slide, 5.65, 3.35, 6.35, 0.45, "Результат", 12, BLUE, True)
    add_text(slide, 5.65, 3.85, 6.4, 1.35, "Large Language Models —\nведущая область", 27, NAVY, True)
    add_text(slide, 5.65, 5.45, 6.25, 0.7, "Trending_Score 95 · Avg_Citations 12.3", 16, INK)

    slide = presentation.slides.add_slide(blank)
    add_title(slide, "Как собраны данные", "Метод", 2)
    stages = [
        (0.8, "01", "scholarly_papers.json", "4 статьи\nPaper_ID · Title · Area", BLUE),
        (4.75, "02", "Точное сопоставление", "Area = Research_Area", CYAN),
        (8.7, "03", "dashboard.html", "Citations = Avg_Citations\nRelevance = Trending_Score", CORAL),
    ]
    for left, num, title, body, accent in stages:
        add_rect(slide, left, 2.05, 3.45, 2.55, WHITE, True, RGBColor(213, 224, 235))
        add_text(slide, left + 0.25, 2.25, 0.55, 0.35, num, 11, accent, True)
        add_text(slide, left + 0.25, 2.83, 2.95, 0.55, title, 17, NAVY, True)
        add_text(slide, left + 0.25, 3.62, 2.95, 0.65, body, 12, INK)
    for left in (4.28, 8.23):
        add_text(slide, left, 3.02, 0.35, 0.4, "→", 23, BLUE, True, PP_ALIGN.CENTER)
    add_rect(slide, 0.8, 5.2, 11.35, 1.0, PALE, True)
    add_text(slide, 1.05, 5.38, 10.85, 0.55, "Проверка: каждая Area должна иметь ровно одну строку на панели; пропуск останавливает анализ.", 13, NAVY, True)
    add_footer(slide)

    slide = presentation.slides.add_slide(blank)
    add_title(slide, "Портфель: 4 статьи × 4 области", "Данные", 3)
    colors = [BLUE, CYAN, CORAL, GOLD]
    for index, (item, accent) in enumerate(zip(result["data_analysis"], colors)):
        row = index // 2
        column = index % 2
        left = 0.8 + column * 6.1
        top = 1.65 + row * 2.35
        add_card(
            slide,
            left,
            top,
            5.7,
            1.85,
            item["Area"],
            f"{item['Title']}\nPaper_ID {item['Paper_ID']}\nCitations {item['Citations']:.1f}  ·  Relevance {item['Relevance_Score']:.0f}",
            accent,
        )
    add_footer(slide)

    slide = presentation.slides.add_slide(blank)
    add_title(slide, "Где сосредоточить внимание", "Выводы", 4)
    ranked = sorted(result["data_analysis"], key=lambda item: item["Relevance_Score"], reverse=True)
    add_text(slide, 0.8, 1.55, 6.5, 0.35, "Relevance_Score", 13, NAVY, True)
    for index, item in enumerate(ranked):
        top = 2.05 + index * 0.9
        add_text(slide, 0.8, top, 2.65, 0.36, item["Area"], 11, INK, True)
        add_rect(slide, 3.55, top + 0.03, 3.25, 0.27, RGBColor(226, 234, 242), False)
        add_rect(slide, 3.55, top + 0.03, 3.25 * item["Relevance_Score"] / 100, 0.27, BLUE if index == 0 else CYAN, False)
        add_text(slide, 6.92, top - 0.02, 0.55, 0.35, f"{item['Relevance_Score']:.0f}", 11, NAVY, True, PP_ALIGN.RIGHT)
        add_text(slide, 7.65, top - 0.02, 1.15, 0.35, f"cit. {item['Citations']:.1f}", 10, MUTED)
    metrics = result["metrics"]
    add_rect(slide, 9.25, 1.65, 3.1, 4.45, NAVY, True)
    add_text(slide, 9.6, 1.95, 2.4, 0.35, "Сводные метрики", 14, WHITE, True)
    for top, value, label in (
        (2.65, str(metrics["Total_Papers"]), "Total_Papers"),
        (3.45, str(metrics["Avg_Citations"]), "Avg_Citations"),
        (4.25, str(metrics["Avg_Relevance"]), "Avg_Relevance"),
    ):
        add_text(slide, 9.6, top, 0.8, 0.5, value, 24, CYAN, True)
        add_text(slide, 10.5, top + 0.08, 1.45, 0.32, label, 10, WHITE)
    add_text(slide, 9.6, 5.2, 2.35, 0.65, f"Top_Area\n{metrics['Top_Area']}", 11, WHITE, True)
    add_footer(slide)

    slide = presentation.slides.add_slide(blank)
    add_title(slide, "Рекомендации и границы", "Действия", 5)
    for index, item in enumerate(result["recommendations"]):
        left = 0.8 + index * 6.05
        accent = CORAL if item["Priority"] == "High" else BLUE
        add_card(slide, left, 1.7, 5.65, 2.15, f"{item['Priority']} · {item['Area']}", item["Action"], accent)
    add_rect(slide, 0.8, 4.4, 11.7, 1.65, PALE, True)
    add_text(slide, 1.1, 4.68, 2.15, 0.35, "Граница вывода", 14, NAVY, True)
    add_text(slide, 3.45, 4.55, 8.6, 0.95, "Gap и направление тренда не рассчитывались: во входе нет сопоставимой внутренней базы и временного ряда.", 14, INK, True)
    add_text(slide, 0.8, 6.35, 11.4, 0.42, "Следующий шаг: обновлять HTML-срез ежеквартально и повторно запускать тот же процессор.", 13, NAVY, True)
    add_footer(slide)

    presentation.save(path)


def solve(input_dir: Path, output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    processor_source = Path(__file__).with_name(PROCESSOR_NAME)
    processor_output = output_dir / PROCESSOR_NAME
    if processor_source.resolve() != processor_output.resolve():
        shutil.copy2(processor_source, processor_output)
    subprocess.run(
        [
            sys.executable,
            str(processor_output),
            "--papers",
            str(input_dir / "scholarly_papers.json"),
            "--dashboard",
            str(input_dir / "dashboard.html"),
            "--output",
            str(output_dir / JSON_NAME),
        ],
        check=True,
        timeout=60,
    )
    result = json.loads((output_dir / JSON_NAME).read_text(encoding="utf-8"))
    write_workbook(output_dir / XLSX_NAME, result)
    write_presentation(output_dir / PPTX_NAME, result)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=Path("/app/data/input_files"))
    parser.add_argument("--output-dir", type=Path, default=Path("/app/output"))
    args = parser.parse_args()
    solve(args.input_dir, args.output_dir)


if __name__ == "__main__":
    main()
