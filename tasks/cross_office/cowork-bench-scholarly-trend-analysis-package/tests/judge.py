#!/usr/bin/env python3
"""学术趋势分析包 v2 连续评分；v1 严格口径保存在 judge_v1.py。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from pptx import Presentation


TASK_ID = "cowork-evo/cowork-bench-scholarly-trend-analysis-package"
RUBRIC_VERSION = "cowork-bench-scholarly-trend-analysis-package-v0.2"
JUDGE_VERSION = "scholarly-trend-criterion-partial-v2.0"
PROCESSOR = "scholarly_trend_processor.py"
RESULT_JSON = "scholarly_trend_results.json"
REPORT = "Trend_Analysis_Report.xlsx"
DECK = "Trend_Analysis_Presentation.pptx"
REQUIRED = (PROCESSOR, RESULT_JSON, REPORT, DECK)
DATA_HEADERS = ["Paper_ID", "Title", "Area", "Citations", "Relevance_Score"]
METRIC_ORDER = ["Total_Papers", "Avg_Citations", "Top_Area", "Avg_Relevance"]
REC_HEADERS = ["Priority", "Action", "Area"]
POSITIVE_WEIGHTS = {
    "R001": 6, "R002": 9, "R003": 9, "R004": 9, "R005": 6,
    "R006": 9, "R007": 8, "R008": 5, "R009": 9, "R010": 5,
    "R011": 6, "R012": 5, "R013": 4,
}
PENALTY_WEIGHTS = {"P001": -12, "P002": -4}


def _load_v1():
    path = Path(__file__).with_name("judge_v1.py")
    spec = importlib.util.spec_from_file_location("scholarly_trend_judge_v1", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v1 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


V1 = _load_v1()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _same(actual: Any, expected: Any) -> bool:
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        return V1.close(actual, expected, 0.011)
    return str(actual or "").strip() == str(expected or "").strip()


def _ratio(findings: Iterable[tuple[bool, str]]) -> tuple[float, list[str], int, int]:
    items = list(findings)
    earned = sum(bool(ok) for ok, _ in items)
    total = len(items)
    return earned / max(1, total), [label for ok, label in items if not ok], earned, total


def _criterion(identifier: str, score: float, failures: list[str], earned: float, total: float) -> dict[str, Any]:
    return {
        "criterion_id": identifier,
        "score": round(max(0.0, min(1.0, score)), 6),
        "evidence": "全部任务专属原子检查通过" if not failures else "未通过：" + "；".join(failures),
        "method": "deterministic",
        "raw": {"earned_units": round(earned, 6), "total_units": round(total, 6)},
    }


def _core(artifact: Path) -> tuple[bool, list[str]]:
    if not artifact.is_dir():
        return False, ["输出目录不存在"]
    errors: list[str] = []
    for name in REQUIRED:
        path = artifact / name
        if not path.is_file() or path.stat().st_size == 0:
            errors.append(f"核心交付缺失或为空：{name}")
            continue
        try:
            if name == PROCESSOR:
                source = path.read_text(encoding="utf-8")
                compile(source, str(path), "exec")
                if len(source.strip()) < 100:
                    raise ValueError("处理脚本过短")
            elif name == RESULT_JSON:
                payload = json.loads(path.read_text(encoding="utf-8"))
                if not isinstance(payload, dict):
                    raise ValueError("JSON 顶层不是对象")
            elif name == REPORT:
                if not zipfile.is_zipfile(path):
                    raise ValueError("不是有效 OOXML ZIP")
                book = load_workbook(path, read_only=True, data_only=False)
                if book.sheetnames != ["Data_Analysis", "Metrics", "Recommendations"]:
                    raise ValueError(f"工作表错误：{book.sheetnames}")
                book.close()
            elif name == DECK:
                if not zipfile.is_zipfile(path):
                    raise ValueError("不是有效 OOXML ZIP")
                deck = Presentation(path)
                if not deck.slides:
                    raise ValueError("PPT 没有页面")
        except Exception as exc:
            errors.append(f"核心交付无法解析：{name}: {type(exc).__name__}: {exc}")
    return not errors, errors


def _fatal(errors: list[str], hashes: list[str]) -> dict[str, Any]:
    identifiers = [*POSITIVE_WEIGHTS, *PENALTY_WEIGHTS]
    return {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes or ["0" * 64],
        "criterion_results": [_criterion(key, 0, errors, 0, 0) for key in identifiers],
        "reward": 0.0, "passed": False, "hurdle_failures": ["R001"], "penalties": [],
    }


def _inputs() -> tuple[Path, Path]:
    root = Path(__file__).resolve().parents[1]
    folder = Path("/app/input_files") if Path("/app/input_files").is_dir() else root / "data" / "input_files"
    return folder / "scholarly_papers.json", folder / "dashboard.html"


def _expected(papers: Path, dashboard: Path) -> dict[str, Any]:
    value = V1.derive_expected(papers, dashboard)
    ranked = sorted(value["data_analysis"], key=lambda item: (-item["Relevance_Score"], -item["Citations"], item["Area"]))
    value["recommendations"] = [
        {"Priority": "High", "Action": "Prioritize quarterly literature review and replication planning", "Area": ranked[0]["Area"]},
        {"Priority": "Medium", "Action": "Track benchmark movement and prepare a focused methods scan", "Area": ranked[1]["Area"]},
    ]
    return value


def _row_findings(actual: Any, expected: list[dict[str, Any]], label: str) -> list[tuple[bool, str]]:
    rows = actual if isinstance(actual, list) else []
    findings = [(isinstance(actual, list), f"{label} 应为数组"), (len(rows) == len(expected), f"{label} 行数错误")]
    for index, wanted in enumerate(expected):
        seen = rows[index] if index < len(rows) and isinstance(rows[index], dict) else {}
        findings.append((list(seen) == DATA_HEADERS, f"{label} 第{index + 1}行字段顺序错误"))
        for key in DATA_HEADERS:
            findings.append((_same(seen.get(key), wanted[key]), f"{label} 第{index + 1}行 {key} 错误"))
    return findings


def _metric_findings(actual: Any, expected: dict[str, Any], label: str) -> list[tuple[bool, str]]:
    value = actual if isinstance(actual, dict) else {}
    findings = [(isinstance(actual, dict), f"{label} 应为对象"), (list(value) == METRIC_ORDER, f"{label} 字段或顺序错误")]
    findings.extend((_same(value.get(key), expected[key]), f"{label}.{key} 错误") for key in METRIC_ORDER)
    return findings


def _rec_findings(actual: Any, expected: list[dict[str, Any]], label: str) -> list[tuple[bool, str]]:
    rows = actual if isinstance(actual, list) else []
    findings = [(isinstance(actual, list), f"{label} 应为数组"), (len(rows) == 2, f"{label} 应恰有2条")]
    for index, wanted in enumerate(expected):
        seen = rows[index] if index < len(rows) and isinstance(rows[index], dict) else {}
        findings.append((list(seen) == REC_HEADERS, f"{label} 第{index + 1}条字段顺序错误"))
        findings.extend((_same(seen.get(key), wanted[key]), f"{label} 第{index + 1}条 {key} 错误") for key in REC_HEADERS)
    return findings


def _bundle_findings(actual: Any, expected: dict[str, Any], label: str) -> list[tuple[bool, str]]:
    payload = actual if isinstance(actual, dict) else {}
    findings = [(list(payload) == ["data_analysis", "metrics", "recommendations"], f"{label} 顶层字段或顺序错误")]
    findings.extend(_row_findings(payload.get("data_analysis"), expected["data_analysis"], f"{label}.data_analysis"))
    findings.extend(_metric_findings(payload.get("metrics"), expected["metrics"], f"{label}.metrics"))
    findings.extend(_rec_findings(payload.get("recommendations"), expected["recommendations"], f"{label}.recommendations"))
    return findings


def _processor(artifact: Path, mutated: bool) -> tuple[float, list[str], int, int, dict[str, Any] | None]:
    papers, dashboard = _inputs()
    with tempfile.TemporaryDirectory(prefix="scholarly-trend-v2-script-") as raw:
        temp = Path(raw)
        if mutated:
            html = dashboard.read_text(encoding="utf-8").replace("<td>10.8</td>", "<td>14.8</td>", 1).replace("<td>82</td>", "<td>86</td>", 1)
            dashboard = temp / "dashboard-mutated.html"
            dashboard.write_text(html, encoding="utf-8")
        expected = _expected(papers, dashboard)
        actual, run_evidence = V1.run_processor(artifact / PROCESSOR, papers, dashboard, temp / "result.json")
        findings = [(actual is not None, f"处理脚本运行失败：{run_evidence}")]
        if actual is not None:
            findings.extend(_bundle_findings(actual, expected, "processor"))
            if mutated:
                rows = actual.get("data_analysis") if isinstance(actual, dict) else []
                cv = next((row for row in rows or [] if isinstance(row, dict) and row.get("Area") == "Computer Vision"), {})
                findings.extend([
                    (_same(cv.get("Citations"), 14.8), "变异 Citations 未传播为14.8"),
                    (_same(cv.get("Relevance_Score"), 86), "变异 Relevance_Score 未传播为86"),
                    (_same((actual.get("metrics") or {}).get("Avg_Citations"), 11), "变异 Avg_Citations 未重算为11"),
                    (_same((actual.get("metrics") or {}).get("Avg_Relevance"), 80), "变异 Avg_Relevance 未重算为80"),
                ])
        score, failures, earned, total = _ratio(findings)
        return score, failures, earned, total, actual


def _excel_rows(artifact: Path, expected: dict[str, Any]) -> tuple[float, list[str], int, int, list[dict[str, Any]]]:
    book = load_workbook(artifact / REPORT, data_only=True)
    sheet = book["Data_Analysis"]
    rows, missing = V1.read_sheet_rows(sheet, DATA_HEADERS)
    names = book.sheetnames
    book.close()
    findings = [(names == ["Data_Analysis", "Metrics", "Recommendations"], f"工作表错误：{names}"), (not missing, f"缺表头：{missing}")]
    findings.extend(_row_findings(rows, expected["data_analysis"], "Excel.Data_Analysis"))
    score, failures, earned, total = _ratio(findings)
    return score, failures, earned, total, rows


def _excel_metrics(artifact: Path, expected: dict[str, Any], recalculated: Path | None) -> tuple[float, list[str], int, int, dict[str, Any]]:
    formula_book = load_workbook(artifact / REPORT, data_only=False)
    sheet = formula_book["Metrics"]
    labels = [str(sheet[f"A{row}"].value or "") for row in range(2, 6)]
    formulas = {f"B{row}": sheet[f"B{row}"].value for row in range(2, 6)}
    normalized = {cell: V1.formula_norm(value) for cell, value in formulas.items()}
    formula_book.close()
    value_path = recalculated or artifact / REPORT
    value_book = load_workbook(value_path, data_only=True)
    value_sheet = value_book["Metrics"]
    values = {str(value_sheet[f"A{row}"].value or ""): value_sheet[f"B{row}"].value for row in range(2, 6)}
    value_book.close()
    findings = [
        (labels == METRIC_ORDER, "Metrics 指标名称或顺序错误"),
        (normalized["B2"].startswith("=") and "COUNTA" in normalized["B2"] and "DATA_ANALYSIS!A2:A5" in normalized["B2"], "B2 总数公式错误"),
        (normalized["B3"].startswith("=") and "AVERAGE" in normalized["B3"] and "DATA_ANALYSIS!D2:D5" in normalized["B3"], "B3 平均引用公式错误"),
        (normalized["B4"].startswith("=") and all(token in normalized["B4"] for token in ("INDEX", "MAX", "DATA_ANALYSIS!C2:C5", "DATA_ANALYSIS!E2:E5")), "B4 最高领域公式错误"),
        (normalized["B5"].startswith("=") and "AVERAGE" in normalized["B5"] and "DATA_ANALYSIS!E2:E5" in normalized["B5"], "B5 平均相关性公式错误"),
        (recalculated is not None, "LibreOffice 未成功重算 Excel"),
    ]
    findings.extend((_same(values.get(key), expected["metrics"][key]), f"Metrics.{key} 重算值错误") for key in METRIC_ORDER)
    score, failures, earned, total = _ratio(findings)
    return score, failures, earned, total, values


def _excel_recommendations(artifact: Path, delivered: dict[str, Any], expected: dict[str, Any]) -> tuple[float, list[str], int, int, list[dict[str, Any]]]:
    book = load_workbook(artifact / REPORT, data_only=True)
    rows, missing = V1.read_sheet_rows(book["Recommendations"], REC_HEADERS)
    book.close()
    findings = [(not missing, f"Recommendations 缺表头：{missing}")]
    findings.extend(_rec_findings(rows, expected["recommendations"], "Excel.Recommendations"))
    json_rows = delivered.get("recommendations") if isinstance(delivered.get("recommendations"), list) else []
    for index in range(2):
        excel_row = rows[index] if index < len(rows) else {}
        json_row = json_rows[index] if index < len(json_rows) and isinstance(json_rows[index], dict) else {}
        findings.extend((_same(excel_row.get(key), json_row.get(key)), f"第{index + 1}条建议 {key} 跨 Excel/JSON 冲突") for key in REC_HEADERS)
    score, failures, earned, total = _ratio(findings)
    return score, failures, earned, total, rows


def _ppt(artifact: Path, expected: dict[str, Any], delivered: dict[str, Any]) -> tuple[float, list[str], int, int, dict[str, Any]]:
    deck, slide_texts, native, pictures, large_pictures, out_of_bounds = V1.ppt_text_and_structure(artifact / DECK)
    blob = "\n".join(slide_texts)
    norm = V1.normalized(blob)
    aspect = deck.slide_width / deck.slide_height
    content_findings: list[tuple[bool, str]] = [
        (4 <= len(deck.slides) <= 8, f"PPT 页数应为4–8，实际{len(deck.slides)}"),
        (abs(aspect - 16 / 9) < 0.03, f"PPT 不是16:9，实际{aspect:.3f}"),
        ("scholarly_papers.json" in norm, "PPT 缺来源 scholarly_papers.json"),
        ("dashboard.html" in norm, "PPT 缺来源 dashboard.html"),
        ("area = research_area" in norm, "PPT 缺精确关联规则 Area = Research_Area"),
        ("avg_citations" in norm, "PPT 缺 Citations 来源字段 Avg_Citations"),
        ("trending_score" in norm, "PPT 缺 Relevance 来源字段 Trending_Score"),
    ]
    for row in expected["data_analysis"]:
        for key in ("Paper_ID", "Title", "Area"):
            content_findings.append((V1.normalized(row[key]) in norm, f"PPT 缺 {row['Paper_ID']} 的 {key}"))
        content_findings.append((f"{row['Citations']:.1f}" in norm, f"PPT 缺 {row['Paper_ID']} Citations={row['Citations']:.1f}"))
        content_findings.append((f"{row['Relevance_Score']:.0f}" in norm, f"PPT 缺 {row['Paper_ID']} Relevance={row['Relevance_Score']:.0f}"))
        citations_value = f"{row['Citations']:.1f}"
        relevance_value = f"{row['Relevance_Score']:.0f}"
        pair_pattern = rf"citations\s+{re.escape(citations_value)}\s+·\s+relevance\s+{re.escape(relevance_value)}"
        content_findings.append((bool(re.search(pair_pattern, norm)), f"PPT 中 {row['Paper_ID']} 的 Citations/Relevance 数值对错误"))
    metric_tokens = {"Total_Papers": "4", "Avg_Citations": "10", "Top_Area": "large language models", "Avg_Relevance": "79"}
    for key, token in metric_tokens.items():
        content_findings.append((V1.normalized(token) in norm, f"PPT 缺指标 {key}={token}"))
    for row in delivered.get("recommendations", []) if isinstance(delivered.get("recommendations"), list) else []:
        content_findings.append((V1.normalized(row.get("Area")) in norm, f"PPT 缺建议领域 {row.get('Area')}"))
    content_score, content_failures, content_earned, content_total = _ratio(content_findings)
    empty_slides = [index for index, text in enumerate(slide_texts, 1) if len(V1.normalized(text)) < 40]
    structure_findings = [
        (native >= 25, f"原生对象少于25，实际{native}"),
        (not large_pictures, f"存在覆盖60%以上页面的大图：{large_pictures}"),
        (not out_of_bounds, f"存在越界对象：{out_of_bounds}"),
        (not empty_slides, f"存在实质空白页：{empty_slides}"),
    ]
    structure_score, structure_failures, structure_earned, structure_total = _ratio(structure_findings)
    return content_score, content_failures, content_earned, content_total, {
        "text": blob, "norm": norm, "slide_count": len(deck.slides),
        "structure": (structure_score, structure_failures, structure_earned, structure_total),
    }


def _office_bundle(artifact: Path) -> dict[str, Any]:
    with tempfile.TemporaryDirectory(prefix="scholarly-trend-v2-office-") as raw:
        recalculated, texts, pages, evidence = V1.libreoffice_prepare(artifact / REPORT, artifact / DECK, Path(raw))
        retained = None
        if recalculated is not None:
            retained = Path(raw) / "recalculated-copy.xlsx"
            shutil.copy2(recalculated, retained)
            # Workbook values must be consumed before the temporary directory closes.
            book = load_workbook(retained, data_only=True)
            metric_values = {str(book["Metrics"][f"A{row}"].value or ""): book["Metrics"][f"B{row}"].value for row in range(2, 6)}
            book.close()
        else:
            metric_values = {}
        return {"recalculated": recalculated is not None, "metric_values": metric_values, "texts": texts, "pages": pages, "evidence": evidence}


def _metrics_with_office(artifact: Path, expected: dict[str, Any], office: dict[str, Any]) -> tuple[float, list[str], int, int, dict[str, Any]]:
    formula_book = load_workbook(artifact / REPORT, data_only=False)
    sheet = formula_book["Metrics"]
    labels = [str(sheet[f"A{row}"].value or "") for row in range(2, 6)]
    formulas = {f"B{row}": sheet[f"B{row}"].value for row in range(2, 6)}
    normalized = {cell: V1.formula_norm(value) for cell, value in formulas.items()}
    formula_book.close()
    findings = [
        (labels == METRIC_ORDER, "Metrics 指标名称或顺序错误"),
        (normalized["B2"].startswith("=") and "COUNTA" in normalized["B2"] and "DATA_ANALYSIS!A2:A5" in normalized["B2"], "B2 总数公式错误"),
        (normalized["B3"].startswith("=") and "AVERAGE" in normalized["B3"] and "DATA_ANALYSIS!D2:D5" in normalized["B3"], "B3 平均引用公式错误"),
        (normalized["B4"].startswith("=") and all(token in normalized["B4"] for token in ("INDEX", "MAX", "DATA_ANALYSIS!C2:C5", "DATA_ANALYSIS!E2:E5")), "B4 最高领域公式错误"),
        (normalized["B5"].startswith("=") and "AVERAGE" in normalized["B5"] and "DATA_ANALYSIS!E2:E5" in normalized["B5"], "B5 平均相关性公式错误"),
        (bool(office["recalculated"]), "LibreOffice 未成功重算 Excel"),
    ]
    findings.extend((_same(office["metric_values"].get(key), expected["metrics"][key]), f"Metrics.{key} 重算值错误") for key in METRIC_ORDER)
    score, failures, earned, total = _ratio(findings)
    return score, failures, earned, total, office["metric_values"]


def _render(office: dict[str, Any], slide_count: int) -> tuple[float, list[str], int, int]:
    xlsx_norm = V1.normalized(office["texts"].get(REPORT, ""))
    ppt_norm = V1.normalized(office["texts"].get(DECK, ""))
    findings = [
        (bool(office["recalculated"]), f"Excel 未重算：{office['evidence']}"),
        (office["pages"].get(REPORT, 0) >= 3, "Excel 渲染少于3页"),
        (office["pages"].get(DECK, 0) == slide_count, "PPT 渲染页数与源文件不一致"),
        (all(token in xlsx_norm for token in ("paper_id", "computer vision", "total_papers", "priority", "action")), "Excel 渲染缺核心文本"),
        (all(token in ppt_norm for token in ("научные", "large language models", "avg_relevance")), "PPT 渲染缺核心文本"),
    ]
    return _ratio(findings)


def _formatting(artifact: Path) -> tuple[float, list[str], int, int]:
    book = load_workbook(artifact / REPORT, data_only=False)
    findings: list[tuple[bool, str]] = []
    for name in ("Data_Analysis", "Metrics", "Recommendations"):
        sheet = book[name]
        headers = list(sheet[1])
        findings.extend([
            (len(sheet.tables) >= 1 or bool(sheet.auto_filter.ref), f"{name} 无原生表或筛选"),
            (str(sheet.freeze_panes or "") == "A2", f"{name} 未冻结首行"),
            (all(bool(cell.font.bold) for cell in headers), f"{name} 表头未全部加粗"),
            (all(cell.fill.fill_type == "solid" for cell in headers), f"{name} 表头未全部填色"),
            (all((sheet.column_dimensions[cell.column_letter].width or 0) >= 9 for cell in headers), f"{name} 存在过窄列"),
        ])
    findings.extend([
        (book["Data_Analysis"]["D2"].number_format == "0.0", "Data_Analysis Citations 数字格式错误"),
        (book["Data_Analysis"]["E2"].number_format == "0", "Data_Analysis Relevance 数字格式错误"),
    ])
    book.close()
    return _ratio(findings)


def _cross(delivered: dict[str, Any], excel_rows: list[dict[str, Any]], metric_values: dict[str, Any], ppt_norm: str) -> tuple[float, list[str], int, int]:
    json_rows = delivered.get("data_analysis") if isinstance(delivered.get("data_analysis"), list) else []
    findings: list[tuple[bool, str]] = [(len(json_rows) == len(excel_rows) == 4, "JSON/Excel 论文行数冲突")]
    for index in range(4):
        json_row = json_rows[index] if index < len(json_rows) and isinstance(json_rows[index], dict) else {}
        excel_row = excel_rows[index] if index < len(excel_rows) else {}
        for key in DATA_HEADERS:
            findings.append((_same(json_row.get(key), excel_row.get(key)), f"第{index + 1}篇 {key} 在 JSON/Excel 冲突"))
        findings.extend([
            (V1.normalized(json_row.get("Paper_ID")) in ppt_norm, f"PPT 缺第{index + 1}篇 Paper_ID"),
            (V1.normalized(json_row.get("Area")) in ppt_norm, f"PPT 缺第{index + 1}篇 Area"),
        ])
    metrics = delivered.get("metrics") if isinstance(delivered.get("metrics"), dict) else {}
    for key in METRIC_ORDER:
        findings.append((_same(metrics.get(key), metric_values.get(key)), f"指标 {key} 在 JSON/Excel 冲突"))
    return _ratio(findings)


def _unsupported(delivered: dict[str, Any], artifact: Path, ppt_text: str) -> tuple[float, list[str], float]:
    book = load_workbook(artifact / REPORT, data_only=True)
    excel_text = "\n".join(str(cell.value or "") for sheet in book.worksheets for row in sheet.iter_rows() for cell in row)
    book.close()
    labels = V1.unsupported_gap_claims(json.dumps(delivered, ensure_ascii=False), excel_text, ppt_text)
    severity = min(1.0, 0.5 + 0.2 * len(labels)) if labels else 0.0
    return severity, labels, 1.0


def _extra(artifact: Path) -> tuple[float, list[str], float]:
    allowed = {*REQUIRED, "solve.py", "solve.sh", "__pycache__"}
    extras = sorted(path.name for path in artifact.iterdir() if path.name not in allowed)
    return (1.0 if extras else 0.0), ([f"存在未要求文件：{extras}"] if extras else []), 1.0


def _safe_positive(call, *args):
    try:
        return call(*args)
    except Exception as exc:
        return 0.0, [f"任务专属检查异常：{type(exc).__name__}: {exc}"], 0, 1


def _safe_penalty(call, *args):
    try:
        return call(*args)
    except Exception as exc:
        return 1.0, [f"严重边界无法核验：{type(exc).__name__}: {exc}"], 1.0


def evaluate(artifact: Path) -> dict[str, Any]:
    ok, errors = _core(artifact)
    hashes = [_sha256(artifact / name) for name in REQUIRED if (artifact / name).is_file()]
    if not ok:
        return _fatal(errors, hashes)

    papers, dashboard = _inputs()
    expected = _expected(papers, dashboard)
    delivered = json.loads((artifact / RESULT_JSON).read_text(encoding="utf-8"))
    results = [_criterion("R001", 1, [], 4, 4)]
    scores: dict[str, float] = {"R001": 1.0}

    canonical = _processor(artifact, False)
    mutated = _processor(artifact, True)
    for key, check in (("R002", canonical), ("R003", mutated)):
        score, failures, earned, total, _ = check
        results.append(_criterion(key, score, failures, earned, total)); scores[key] = score

    json_rows = _ratio(_row_findings(delivered.get("data_analysis"), expected["data_analysis"], "JSON.data_analysis"))
    json_summary = _ratio([*_metric_findings(delivered.get("metrics"), expected["metrics"], "JSON.metrics"), *_rec_findings(delivered.get("recommendations"), expected["recommendations"], "JSON.recommendations")])
    for key, check in (("R004", json_rows), ("R005", json_summary)):
        score, failures, earned, total = check
        results.append(_criterion(key, score, failures, earned, total)); scores[key] = score

    office = _office_bundle(artifact)
    excel_rows_check = _safe_positive(_excel_rows, artifact, expected)
    metric_check = _safe_positive(_metrics_with_office, artifact, expected, office)
    rec_check = _safe_positive(_excel_recommendations, artifact, delivered, expected)
    ppt_check = _safe_positive(_ppt, artifact, expected, delivered)
    excel_score, excel_failures, excel_earned, excel_total, excel_rows = excel_rows_check
    metric_score, metric_failures, metric_earned, metric_total, metric_values = metric_check
    rec_score, rec_failures, rec_earned, rec_total, _ = rec_check
    ppt_score, ppt_failures, ppt_earned, ppt_total, ppt_data = ppt_check
    structure_check = ppt_data.get("structure", (0.0, ["PPT 结构无法读取"], 0, 1))
    cross_check = _safe_positive(_cross, delivered, excel_rows, metric_values, ppt_data.get("norm", ""))
    render_check = _safe_positive(_render, office, int(ppt_data.get("slide_count", 0)))
    formatting_check = _safe_positive(_formatting, artifact)
    checks = {
        "R006": (excel_score, excel_failures, excel_earned, excel_total),
        "R007": (metric_score, metric_failures, metric_earned, metric_total),
        "R008": (rec_score, rec_failures, rec_earned, rec_total),
        "R009": (ppt_score, ppt_failures, ppt_earned, ppt_total),
        "R010": structure_check,
        "R011": cross_check,
        "R012": render_check,
        "R013": formatting_check,
    }
    for key, (score, failures, earned, total) in checks.items():
        results.append(_criterion(key, score, failures, earned, total)); scores[key] = score

    penalties = {
        "P001": _safe_penalty(_unsupported, delivered, artifact, ppt_data.get("text", "")),
        "P002": _safe_penalty(_extra, artifact),
    }
    for key, (severity, labels, total) in penalties.items():
        results.append(_criterion(key, severity, labels, severity * total, total)); scores[key] = severity

    total_weight = sum(POSITIVE_WEIGHTS.values())
    positive = sum(POSITIVE_WEIGHTS[key] * scores[key] for key in POSITIVE_WEIGHTS)
    deduction = sum(abs(PENALTY_WEIGHTS[key]) * scores[key] for key in PENALTY_WEIGHTS)
    reward = round(max(0.0, min(1.0, (positive - deduction) / total_weight)), 6)
    return {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes, "criterion_results": results,
        "reward": reward, "passed": reward >= 0.8, "hurdle_failures": [],
        "penalties": [{"id": key, "amount": round(abs(PENALTY_WEIGHTS[key]) * scores[key] / total_weight, 6)} for key in PENALTY_WEIGHTS if scores[key] > 0],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    parser.add_argument("--scoring-version", choices=("v1", "v2"), default="v2")
    args = parser.parse_args()
    try:
        result = V1.evaluate(args.artifact_dir) if args.scoring_version == "v1" else evaluate(args.artifact_dir)
    except Exception as exc:
        result = _fatal([f"Judge 异常：{type(exc).__name__}: {exc}"], [])
    rendered = json.dumps(result, ensure_ascii=False, indent=2)
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(rendered + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{float(result['reward']):.6f}\n", encoding="utf-8")
    print(rendered)


if __name__ == "__main__":
    main()
