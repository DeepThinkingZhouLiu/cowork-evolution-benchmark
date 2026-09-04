#!/usr/bin/env python3
"""学术趋势分析包的任务专属 v2 致命与连续部分分回归。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution"
INPUTS = ROOT / "environment" / "input_files"
REQUIRED = ("scholarly_trend_processor.py", "scholarly_trend_results.json", "Trend_Analysis_Report.xlsx", "Trend_Analysis_Presentation.pptx")
Mutator = Callable[[Path], None]


def _load_judge():
    path = Path(__file__).with_name("judge.py")
    spec = importlib.util.spec_from_file_location("scholarly_trend_v2_judge_for_negatives", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v2 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


JUDGE = _load_judge()


def _copy(target: Path) -> None:
    target.mkdir(parents=True, exist_ok=True)
    for name in (*REQUIRED, "solve.py", "solve.sh"):
        source = GOLDEN / name
        if source.is_file():
            shutil.copy2(source, target / name)


def _json(mutator: Callable[[dict], None]) -> Mutator:
    def apply(root: Path) -> None:
        path = root / "scholarly_trend_results.json"
        payload = json.loads(path.read_text(encoding="utf-8"))
        mutator(payload)
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return apply


def _excel(sheet: str, cell: str, value: object) -> Mutator:
    def apply(root: Path) -> None:
        path = root / "Trend_Analysis_Report.xlsx"
        book = load_workbook(path)
        book[sheet][cell] = value
        book.save(path)
        book.close()
    return apply


def _processor_mapping(root: Path) -> None:
    path = root / "scholarly_trend_processor.py"
    text = path.read_text(encoding="utf-8")
    path.write_text(text.replace('area = str(paper["Area"])', 'area = "Federated Learning" if paper["Area"] == "Computer Vision" else str(paper["Area"])'), encoding="utf-8")


def _processor_hardcoded(root: Path) -> None:
    payload = json.loads((root / "scholarly_trend_results.json").read_text(encoding="utf-8"))
    script = (
        "#!/usr/bin/env python3\nimport argparse,json\nfrom pathlib import Path\n"
        "p=argparse.ArgumentParser();p.add_argument('--papers');p.add_argument('--dashboard');p.add_argument('--output',type=Path);a=p.parse_args()\n"
        f"payload={payload!r}\n"
        "a.output.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+'\\n',encoding='utf-8')\n"
    )
    (root / "scholarly_trend_processor.py").write_text(script, encoding="utf-8")


def _format_header(root: Path) -> None:
    path = root / "Trend_Analysis_Report.xlsx"
    book = load_workbook(path)
    book["Data_Analysis"]["A1"].fill = PatternFill(fill_type=None)
    book.save(path)
    book.close()


def _ppt_replace(old: str, new: str) -> Mutator:
    def apply(root: Path) -> None:
        path = root / "Trend_Analysis_Presentation.pptx"
        deck = Presentation(path)
        replacements = 0
        for slide in deck.slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and old in shape.text:
                    shape.text = shape.text.replace(old, new)
                    replacements += 1
        if not replacements:
            raise AssertionError(f"PPT 未找到待替换文本：{old}")
        deck.save(path)
    return apply


def _ppt_one_object_out_of_bounds(root: Path) -> None:
    path = root / "Trend_Analysis_Presentation.pptx"
    deck = Presentation(path)
    shape = deck.slides[1].shapes[0]
    shape.left = -shape.width
    deck.save(path)


def _ppt_unsupported_gap(root: Path) -> None:
    path = root / "Trend_Analysis_Presentation.pptx"
    deck = Presentation(path)
    shape = deck.slides[-1].shapes.add_textbox(100000, 100000, 3500000, 400000)
    shape.text = "Gap: 27% and Large Language Models improving"
    deck.save(path)


def _moderate_json_conflict(root: Path) -> None:
    path = root / "scholarly_trend_results.json"
    value = json.loads(path.read_text(encoding="utf-8"))
    for index, row in enumerate(value["data_analysis"], 1):
        row["Citations"] = float(row["Citations"]) + index
        row["Relevance_Score"] = float(row["Relevance_Score"]) - index
    value["metrics"] = {"Total_Papers": 5, "Avg_Citations": 13, "Top_Area": "Computer Vision", "Avg_Relevance": 70}
    value["recommendations"][0]["Area"] = "Federated Learning"
    value["recommendations"][1]["Area"] = "Reinforcement Learning"
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _severe_valid_incomplete(root: Path) -> None:
    json_path = root / "scholarly_trend_results.json"
    value = json.loads(json_path.read_text(encoding="utf-8"))
    value["data_analysis"] = []
    value["metrics"] = {}
    value["recommendations"] = []
    json_path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    excel_path = root / "Trend_Analysis_Report.xlsx"
    book = load_workbook(excel_path)
    for sheet_name in ("Data_Analysis", "Recommendations"):
        sheet = book[sheet_name]
        if sheet.max_row > 1:
            sheet.delete_rows(2, sheet.max_row - 1)
    for row in range(2, 6):
        book["Metrics"][f"B{row}"] = 0
    book.save(excel_path)
    book.close()

    ppt_path = root / "Trend_Analysis_Presentation.pptx"
    deck = Presentation(ppt_path)
    for slide in list(deck.slides)[1:]:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                shape.text = ""
    deck.save(ppt_path)


def _input_only(root: Path) -> None:
    for path in list(root.iterdir()):
        if path.is_file():
            path.unlink()
    for source in INPUTS.iterdir():
        if source.is_file():
            shutil.copy2(source, root / source.name)


PARTIALS: dict[str, Mutator] = {
    "json_one_paper_id_wrong": _json(lambda value: value["data_analysis"][0].update({"Paper_ID": "2302.99999"})),
    "json_one_citation_wrong": _json(lambda value: value["data_analysis"][1].update({"Citations": 7.9})),
    "json_avg_relevance_wrong": _json(lambda value: value["metrics"].update({"Avg_Relevance": 80})),
    "json_top_area_wrong": _json(lambda value: value["metrics"].update({"Top_Area": "Computer Vision"})),
    "json_one_recommendation_priority_wrong": _json(lambda value: value["recommendations"][1].update({"Priority": "Low"})),
    "json_one_recommendation_action_wrong": _json(lambda value: value["recommendations"][0].update({"Action": "Read more"})),
    "processor_one_area_mapping_wrong": _processor_mapping,
    "processor_hardcoded_against_mutation": _processor_hardcoded,
    "excel_one_paper_id_wrong": _excel("Data_Analysis", "A2", "2302.99999"),
    "excel_one_citation_wrong": _excel("Data_Analysis", "D3", 7.9),
    "excel_total_formula_flattened": _excel("Metrics", "B2", 4),
    "excel_top_area_formula_flattened": _excel("Metrics", "B4", "Large Language Models"),
    "excel_one_recommendation_priority_wrong": _excel("Recommendations", "A3", "Low"),
    "excel_one_header_fill_missing": _format_header,
    "ppt_one_paper_id_missing": _ppt_replace("2302.05678", "2302.99999"),
    "ppt_one_title_missing": _ppt_replace("Federated Learning Survey", "Federated Study"),
    "ppt_dashboard_source_missing": _ppt_replace("dashboard.html", "web dashboard"),
    "ppt_mapping_rule_missing": _ppt_replace("Area = Research_Area", "Area mapping"),
    "ppt_one_relevance_wrong": _ppt_replace("Relevance 68", "Relevance 69"),
    "ppt_one_object_out_of_bounds": _ppt_one_object_out_of_bounds,
    "ppt_unsupported_gap_claim": _ppt_unsupported_gap,
    "moderate_json_cross_artifact_conflict": _moderate_json_conflict,
    "severe_valid_but_incomplete_package": _severe_valid_incomplete,
    "unexpected_debug_file": lambda root: (root / "debug-notes.txt").write_text("unrequested\n", encoding="utf-8"),
}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    golden_office = JUDGE._office_bundle(GOLDEN)
    golden_result = JUDGE.evaluate(GOLDEN)
    assert golden_result["reward"] == 1.0, golden_result
    JUDGE._office_bundle = lambda _artifact: golden_office
    records = []
    with tempfile.TemporaryDirectory(prefix="scholarly-trend-v2-negatives-") as raw:
        base = Path(raw)
        for name in REQUIRED:
            for mode in ("missing", "empty", "corrupt"):
                case = base / f"fatal-{mode}-{name}"
                _copy(case)
                target = case / name
                if mode == "missing":
                    target.unlink()
                elif mode == "empty":
                    target.write_bytes(b"")
                else:
                    target.write_bytes(b"not-a-valid-task-artifact")
                reward = float(JUDGE.evaluate(case)["reward"])
                assert reward == 0, (mode, name, reward)
                records.append({"case": f"{mode}:{name}", "layer": "fatal_zero", "reward": reward})

        empty = base / "fatal-empty-output"
        empty.mkdir()
        reward = float(JUDGE.evaluate(empty)["reward"])
        assert reward == 0
        records.append({"case": "empty-output", "layer": "fatal_zero", "reward": reward})

        input_only = base / "fatal-input-only"
        _copy(input_only)
        _input_only(input_only)
        reward = float(JUDGE.evaluate(input_only)["reward"])
        assert reward == 0
        records.append({"case": "unmodified-input-only", "layer": "fatal_zero", "reward": reward})

        for name, mutator in PARTIALS.items():
            case = base / f"partial-{name}"
            _copy(case)
            mutator(case)
            reward = float(JUDGE.evaluate(case)["reward"])
            assert 0 < reward < 1, (name, reward)
            records.append({"case": name, "layer": "partial", "reward": reward})

    partial = [row["reward"] for row in records if row["layer"] == "partial"]
    report = {
        "total": len(records),
        "fatal_zero": sum(row["layer"] == "fatal_zero" for row in records),
        "partial_continuous": len(partial),
        "partial_reward_range": [min(partial), max(partial)],
        "cases": records,
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
