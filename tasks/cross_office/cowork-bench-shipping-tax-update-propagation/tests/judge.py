#!/usr/bin/env python3
"""配送与税率更新传播任务的 v2 连续评分；v1 保存在 judge_v1.py。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from openpyxl import load_workbook


TASK_ID = "cowork-evo/cowork-bench-shipping-tax-update-propagation"
RUBRIC_VERSION = "cowork-bench-shipping-tax-update-propagation-v0.2"
JUDGE_VERSION = "shipping-tax-atomic-propagation-v2.0"
PROCESSOR = "shipping_tax_update_processor.py"
RESULT_JSON = "shipping_tax_results.json"
REPORT = "Shipping_Tax_Analysis.xlsx"
WORD = "Operations_Report.docx"
EMAIL = "operations_email.txt"
REQUIRED = (PROCESSOR, RESULT_JSON, REPORT, WORD, EMAIL)
GENERATED = (RESULT_JSON, REPORT, WORD, EMAIL)
POSITIVE_WEIGHTS = {"R001":6,"R002":10,"R003":12,"R004":9,"R005":7,"R006":8,"R007":6,"R008":9,"R009":7,"R010":6,"R011":4,"R012":5}
PENALTY_WEIGHTS = {"P001":-12,"P002":-4}


def _load_v1():
    path = Path(__file__).with_name("judge_v1.py")
    spec = importlib.util.spec_from_file_location("shipping_tax_judge_v1", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v1 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


V1 = _load_v1()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _same(actual: Any, expected: Any) -> bool:
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        return V1.close_number(actual, expected)
    if expected is None:
        return actual in (None, "")
    return actual == expected


def _ratio(findings: Iterable[tuple[bool, str]]) -> tuple[float, list[str], int, int]:
    items = list(findings)
    earned = sum(bool(ok) for ok, _ in items)
    total = len(items)
    return earned / max(1, total), [label for ok, label in items if not ok], earned, total


def _criterion(identifier: str, score: float, failures: list[str], earned: float, total: float) -> dict[str, Any]:
    return {
        "criterion_id":identifier,
        "score":round(max(0.0,min(1.0,score)),6),
        "evidence":"全部任务专属原子检查通过" if not failures else "未通过："+"；".join(failures),
        "method":"deterministic",
        "raw":{"earned_units":round(earned,6),"total_units":round(total,6)},
    }


def _core(artifact: Path) -> tuple[bool,list[str]]:
    if not artifact.is_dir():
        return False,["输出目录不存在"]
    errors=[]
    for name in REQUIRED:
        path=artifact/name
        if not path.is_file() or path.stat().st_size==0:
            errors.append(f"核心交付缺失或为空：{name}");continue
        try:
            if name==PROCESSOR:
                text=path.read_text(encoding="utf-8");compile(text,str(path),"exec")
                if len(text.strip())<100:raise ValueError("处理脚本过短")
            elif name==RESULT_JSON:
                if not isinstance(json.loads(path.read_text(encoding="utf-8")),dict):raise ValueError("JSON 顶层不是对象")
            elif name==REPORT:
                if not zipfile.is_zipfile(path):raise ValueError("不是有效 OOXML ZIP")
                book=load_workbook(path,read_only=True,data_only=False)
                if book.sheetnames != ["Shipping Zones","Tax Rates"]:raise ValueError(f"工作表错误：{book.sheetnames}")
                book.close()
            elif name==WORD:
                if not zipfile.is_zipfile(path):raise ValueError("不是有效 OOXML ZIP")
                doc=Document(path)
                if not doc.paragraphs or not doc.tables:raise ValueError("Word 无正文或原生表格")
            else:
                if len(path.read_text(encoding="utf-8").strip())<40:raise ValueError("邮件内容过短")
        except Exception as exc:
            errors.append(f"核心交付无法解析：{name}: {type(exc).__name__}: {exc}")
    return not errors,errors


def _fatal(errors:list[str],hashes:list[str])->dict[str,Any]:
    return {"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":hashes or ["0"*64],"criterion_results":[_criterion(key,0,errors,0,0) for key in [*POSITIVE_WEIGHTS,*PENALTY_WEIGHTS]],"reward":0.0,"passed":False,"hurdle_failures":["R001"],"penalties":[]}


def _input_dir()->Path:
    root=Path(__file__).resolve().parents[1]
    return Path("/app/input_files") if Path("/app/input_files").is_dir() else root/"data"/"input_files"


def _atomic(actual:Any,expected:Any,path:str)->list[tuple[bool,str]]:
    findings=[]
    if isinstance(expected,dict):
        source=actual if isinstance(actual,dict) else {}
        findings.extend([(isinstance(actual,dict),f"{path} 应为对象"),(list(source)==list(expected),f"{path} 键或顺序错误")])
        for key,value in expected.items():findings.extend(_atomic(source.get(key),value,f"{path}.{key}"))
    elif isinstance(expected,list):
        rows=actual if isinstance(actual,list) else []
        findings.extend([(isinstance(actual,list),f"{path} 应为数组"),(len(rows)==len(expected),f"{path} 数量错误")])
        for index,value in enumerate(expected):findings.extend(_atomic(rows[index] if index<len(rows) else None,value,f"{path}[{index}]") )
    else:
        findings.append((_same(actual,expected),f"{path} 错误：实际={actual!r}，期望={expected!r}"))
    return findings


def _read_excel_atomic(artifact:Path,truth:dict[str,Any])->dict[str,Any]:
    book=load_workbook(artifact/REPORT,data_only=False)
    shipping=book["Shipping Zones"];taxes=book["Tax Rates"]
    shipping_headers=[shipping.cell(1,c).value for c in range(1,5)]
    tax_headers=[taxes.cell(1,c).value for c in range(1,6)]
    expected_shipping=V1.expected_shipping_excel(truth)
    actual_shipping=[];formula_findings=[]
    for row in range(2,shipping.max_row+1):
        if all(shipping.cell(row,c).value in (None,"") for c in range(1,5)):continue
        methods=str(shipping.cell(row,4).value or "")
        actual_shipping.append({"Zone_Name":shipping.cell(row,1).value,"Methods_Count":0 if not methods else len(methods.split("; ")),"Regions":shipping.cell(row,3).value or "","Methods":methods})
        formula=shipping.cell(row,2).value
        formula_findings.append((V1.dynamic_count_formula(formula,row),f"Shipping Zones 第{row}行 Methods_Count 未动态引用 D{row}"))
    expected_tax=V1.expected_tax_excel(truth);actual_tax=[]
    for row in range(2,taxes.max_row+1):
        if all(taxes.cell(row,c).value in (None,"") for c in range(1,6)):continue
        values=[taxes.cell(row,c).value for c in range(1,6)]
        actual_tax.append({"Tax_Class":values[0] or "","State":values[1] or "","Rate":values[2],"Compound":values[3],"Shipping":values[4]})
    shipping_findings=[(shipping_headers==["Zone_Name","Methods_Count","Regions","Methods"],"Shipping Zones 表头错误")]
    shipping_findings.extend(_atomic(actual_shipping,expected_shipping,"Excel.Shipping Zones"))
    tax_findings=[(tax_headers==["Tax_Class","State","Rate","Compound","Shipping"],"Tax Rates 表头错误")]
    tax_findings.extend(_atomic(actual_tax,expected_tax,"Excel.Tax Rates"))
    formatting=[]
    for sheet in (shipping,taxes):
        headers=list(sheet[1])
        formatting.extend([(len(sheet.tables)>=1 or bool(sheet.auto_filter.ref),f"{sheet.title} 无原生表或筛选"),(str(sheet.freeze_panes or "")=="A2",f"{sheet.title} 未冻结首行"),(all(bool(cell.font.bold) for cell in headers),f"{sheet.title} 表头未全部加粗"),(all(cell.fill.fill_type=="solid" for cell in headers),f"{sheet.title} 表头未全部填色"),(all((sheet.column_dimensions[cell.column_letter].width or 0)>=9 for cell in headers),f"{sheet.title} 存在过窄列")])
    formatting.extend([(len(shipping._images)==0,"Shipping Zones 以图片替代表格"),(len(taxes._images)==0,"Tax Rates 以图片替代表格"),(taxes["C2"].number_format=='0.0000"%"',"Tax Rates Rate 数字格式错误")])
    book.close()
    return {"shipping":_ratio(shipping_findings),"tax":_ratio(tax_findings),"formula":_ratio(formula_findings),"format":formatting,"shipping_rows":actual_shipping,"tax_rows":actual_tax}


def _word_atomic(artifact:Path,truth:dict[str,Any])->tuple[float,list[str],int,int,str]:
    text,headings,tables,media=V1.extract_docx(artifact/WORD)
    findings=[]
    heading_norm={item.casefold() for item in headings}
    findings.extend((heading.casefold() in heading_norm,f"Word 缺章节 {heading}") for heading in V1.REQUIRED_HEADINGS)
    summary=truth["summary"]
    for label,key in (("Shipping zones","total_zones"),("Enabled methods","total_enabled_methods"),("Tax classes","tax_classes"),("Configured tax rates","configured_tax_rates"),("Displayed tax rows","displayed_tax_rows")):
        findings.append((V1.contains_count_near(text,label,summary[key]),f"Word 汇总缺 {label}={summary[key]}"))
    findings.append((f"{V1.rate_text(summary['highest_configured_rate'])}%" in text,"Word 最高税率错误"))
    for zone in truth["shipping_zones"]:
        rows=V1.find_table_rows(tables,zone["zone_name"])
        regions="; ".join(f"{x['type']}:{x['code']}" for x in zone["locations"])
        methods="; ".join(x["title"] for x in zone["enabled_methods"])
        findings.extend([(bool(rows),f"Word 缺配送区 {zone['zone_name']}"),(any(len(row)>=4 and row[1].strip()==str(zone["methods_count"]) for row in rows),f"Word {zone['zone_name']} 方法数错误"),(any(len(row)>=4 and row[2].strip()==regions for row in rows),f"Word {zone['zone_name']} 地域错误"),(any(len(row)>=4 and row[3].strip()==methods for row in rows),f"Word {zone['zone_name']} 方法标题错误")])
    for tax_class in truth["tax_classes"]:
        rows=V1.find_table_rows(tables,tax_class["name"])
        if not tax_class["rates"]:
            findings.extend([(bool(rows),f"Word 缺税类 {tax_class['name']}"),(any(len(row)>=2 and row[1].strip().casefold()=="no configured rate" for row in rows),"Word Zero Rate 未声明无配置")]);continue
        for rate in tax_class["rates"]:
            scope=" / ".join(v for v in (rate["country"],rate["state"],rate["postcode"],rate["city"]) if v)
            matching=[row for row in rows if len(row)>=6 and row[1].strip()==rate["rate_name"]]
            findings.extend([(bool(matching),f"Word 缺税率 {rate['rate_id']}/{rate['rate_name']}"),(any(row[2].strip()==scope for row in matching),f"Word 税率 {rate['rate_id']} 范围错误"),(any(row[3].strip()==f"{V1.rate_text(rate['rate'])}%" for row in matching),f"Word 税率 {rate['rate_id']} 数值错误"),(any(row[4].strip().casefold()==str(rate["compound"]).lower() for row in matching),f"Word 税率 {rate['rate_id']} Compound错误"),(any(row[5].strip().casefold()==str(rate["shipping"]).lower() for row in matching),f"Word 税率 {rate['rate_id']} Shipping错误")])
    findings.extend([(len(tables)>=3,"Word 原生表格少于3张"),(media==0,"Word 用图片替代关键内容"),("No order volume, revenue, customer demand".casefold() in text.casefold(),"Word 缺来源边界声明")])
    score,failures,earned,total=_ratio(findings)
    return score,failures,earned,total,text


def _email_atomic(artifact:Path,truth:dict[str,Any])->tuple[float,list[str],int,int,str]:
    text=(artifact/EMAIL).read_text(encoding="utf-8")
    findings=[(bool(re.search(r"^To:\s*operations@store\.com\s*$",text,re.I|re.M)),"邮件收件人错误"),(bool(re.search(r"^Subject:\s*Shipping and Tax Configuration Report\s*$",text,re.I|re.M)),"邮件主题错误")]
    summary=truth["summary"]
    for label,key in (("shipping zones","total_zones"),("enabled shipping methods","total_enabled_methods"),("tax classes","tax_classes"),("configured tax-rate records","configured_tax_rates")):
        findings.append((V1.contains_count_near(text,label,summary[key],70),f"邮件缺 {label}={summary[key]}"))
    findings.append((f"{V1.rate_text(summary['highest_configured_rate'])}%" in text,"邮件最高税率错误"))
    for zone in truth["shipping_zones"]:
        line=next((line for line in text.splitlines() if zone["zone_name"] in line),"")
        findings.extend([(bool(line),f"邮件缺配送区 {zone['zone_name']}"),(bool(re.search(rf"\b{zone['methods_count']}\b",line)),f"邮件 {zone['zone_name']} 方法数错误")])
    for tax_class in truth["tax_classes"]:
        line=next((line for line in text.splitlines() if tax_class["name"] in line),"")
        findings.extend([(bool(line),f"邮件缺税类 {tax_class['name']}"),(bool(re.search(rf"\b{len(tax_class['rates'])}\b",line)),f"邮件 {tax_class['name']} 记录数错误")])
        for rate in tax_class["rates"]:findings.append((f"{V1.rate_text(rate['rate'])}%" in line,f"邮件缺税率 {rate['rate_id']}={V1.rate_text(rate['rate'])}%"))
        if not tax_class["rates"]:findings.append(("none" in line.casefold(),"邮件 Zero Rate 未写 none"))
    findings.extend([("Recommended checks:" in text,"邮件缺建议"),("draft was not sent" in text,"邮件未声明离线草稿")])
    score,failures,earned,total=_ratio(findings)
    return score,failures,earned,total,text


def _run_processor(script:Path,input_dir:Path,truth:dict[str,Any])->tuple[float,list[str],int,int]:
    with tempfile.TemporaryDirectory(prefix="shipping-tax-v2-replay-") as raw:
        output=Path(raw)/"output"
        completed=subprocess.run([sys.executable,str(script.resolve()),"--input-dir",str(input_dir.resolve()),"--output-dir",str(output)],capture_output=True,text=True,timeout=120,check=False)
        findings=[(completed.returncode==0,f"processor 退出码={completed.returncode}")]
        findings.extend(((output/name).is_file(),f"processor 未生成 {name}") for name in GENERATED)
        if completed.returncode==0 and all((output/name).is_file() for name in GENERATED):
            try:findings.extend(_atomic(json.loads((output/RESULT_JSON).read_text(encoding="utf-8")),truth,"processor.JSON"))
            except Exception as exc:findings.append((False,f"processor JSON不可读：{exc}"))
            try:
                excel=_read_excel_atomic(output,truth)
                for key in ("shipping","tax","formula"):
                    score,failures,earned,total=excel[key]
                    findings.extend((True,f"processor Excel {key} 原子检查通过 {index+1}") for index in range(earned))
                    findings.extend((False,f"processor Excel {key}：{label}") for label in failures)
            except Exception as exc:findings.append((False,f"processor Excel不可读：{exc}"))
            try:
                score,failures,earned,total,_=_word_atomic(output,truth)
                findings.extend((True,f"processor Word 原子检查通过 {index+1}") for index in range(earned))
                findings.extend((False,f"processor Word：{label}") for label in failures)
            except Exception as exc:findings.append((False,f"processor Word不可读：{exc}"))
            try:
                score,failures,earned,total,_=_email_atomic(output,truth)
                findings.extend((True,f"processor Email 原子检查通过 {index+1}") for index in range(earned))
                findings.extend((False,f"processor Email：{label}") for label in failures)
            except Exception as exc:findings.append((False,f"processor Email不可读：{exc}"))
        return _ratio(findings)


def _mutation(artifact:Path,input_dir:Path)->tuple[float,list[str],int,int]:
    with tempfile.TemporaryDirectory(prefix="shipping-tax-v2-mutation-") as raw:
        base=Path(raw);mutated=base/"inputs";V1.build_mutated_inputs(input_dir,mutated);output=base/"output"
        completed=subprocess.run([sys.executable,str((artifact/PROCESSOR).resolve()),"--input-dir",str(mutated),"--output-dir",str(output)],capture_output=True,text=True,timeout=120,check=False)
        findings=[(completed.returncode==0,f"变异回放退出码={completed.returncode}")]
        findings.extend(((output/name).is_file(),f"变异回放缺 {name}") for name in GENERATED)
        if completed.returncode==0 and all((output/name).is_file() for name in GENERATED):
            try:
                payload=json.loads((output/RESULT_JSON).read_text(encoding="utf-8"));rate=next(r for c in payload["tax_classes"] if c["slug"]=="reduced-rate" for r in c["rates"] if r["rate_id"]==7)["rate"]
                findings.extend([(V1.close_number(rate,7),"JSON 未传播 Reduced Rate=7"),(not V1.close_number(rate,5),"JSON 残留 Reduced Rate=5")])
            except Exception as exc:findings.append((False,f"JSON 变异检查失败：{exc}"))
            try:
                book=load_workbook(output/REPORT,data_only=False);rows=[row for row in book["Tax Rates"].iter_rows(min_row=2,values_only=True) if row[0]=="Reduced Rate"];book.close()
                findings.extend([(len(rows)==1,"Excel Reduced Rate 行数错误"),(bool(rows) and V1.close_number(rows[0][2],7),"Excel 未传播 Reduced Rate=7")])
            except Exception as exc:findings.append((False,f"Excel 变异检查失败：{exc}"))
            try:
                text,_,tables,_=V1.extract_docx(output/WORD);rows=V1.find_table_rows(tables,"Reduced Rate")
                findings.extend([(any(len(row)>=4 and row[3].strip()=="7%" for row in rows),"Word 未传播 Reduced Rate=7%"),(not any(len(row)>=4 and row[3].strip()=="5%" for row in rows),"Word Reduced Rate 仍为5%")])
            except Exception as exc:findings.append((False,f"Word 变异检查失败：{exc}"))
            try:
                text=(output/EMAIL).read_text(encoding="utf-8");line=next((x for x in text.splitlines() if "Reduced Rate" in x),"")
                findings.extend([("7%" in line,"邮件未传播 Reduced Rate=7%"),("5%" not in line,"邮件 Reduced Rate 仍为5%")])
            except Exception as exc:findings.append((False,f"邮件变异检查失败：{exc}"))
        return _ratio(findings)


def _cross(truth:dict[str,Any],payload:dict[str,Any],excel:dict[str,Any],word_text:str,email_text:str)->tuple[float,list[str],int,int]:
    findings=[];summary=truth["summary"]
    for zone in truth["shipping_zones"]:
        count=zone["methods_count"];name=zone["zone_name"]
        json_zone=next((x for x in payload.get("shipping_zones",[]) if isinstance(x,dict) and x.get("zone_name")==name),{})
        excel_zone=next((x for x in excel["shipping_rows"] if x.get("Zone_Name")==name),{})
        findings.extend([(_same(json_zone.get("methods_count"),count),f"JSON {name} 方法数冲突"),(_same(excel_zone.get("Methods_Count"),count),f"Excel {name} 方法数冲突"),(V1.contains_count_near(word_text,name,count,120),f"Word {name} 方法数冲突"),(bool(re.search(re.escape(name)+rf"[^\n]*\b{count}\b",email_text)),f"邮件 {name} 方法数冲突")])
    for rate in [r for c in truth["tax_classes"] for r in c["rates"]]:
        token=f"{V1.rate_text(rate['rate'])}%"
        findings.extend([(token in word_text,f"Word 缺税率 {rate['rate_id']}"),(token in email_text,f"邮件缺税率 {rate['rate_id']}")])
    findings.extend([(_same((payload.get("summary") or {}).get("highest_configured_rate"),summary["highest_configured_rate"]),"JSON 最高税率冲突"),(f"{V1.rate_text(summary['highest_configured_rate'])}%" in word_text,"Word 最高税率冲突"),(f"{V1.rate_text(summary['highest_configured_rate'])}%" in email_text,"邮件最高税率冲突")])
    return _ratio(findings)


def _render(artifact:Path)->tuple[float,list[str],int,int]:
    with tempfile.TemporaryDirectory(prefix="shipping-tax-v2-render-") as raw:rendered,log=V1.render_office([artifact/REPORT,artifact/WORD],Path(raw))
    x=rendered.get(REPORT,{});w=rendered.get(WORD,{})
    return _ratio([(x.get("pages",0)>=2,f"Excel 渲染少于2页：{log}"),(w.get("pages",0)>=1,f"Word 渲染失败：{log}"),(all(t in x.get("text","") for t in ("Zone_Name","Methods_Count","Tax_Class","Reduced Rate")),"Excel 渲染缺核心文字"),(all(t in w.get("text","") for t in ("Executive Summary","Shipping Coverage","Tax Structure","Recommendations","Source Boundary")),"Word 渲染缺核心章节")])


def _unsupported(word_text:str,email_text:str,truth:dict[str,Any])->tuple[float,list[str],float]:
    labels=V1.unsupported_claims(word_text+"\n"+email_text,truth);severity=min(1.0,0.45+0.15*len(labels)) if labels else 0.0;return severity,labels,1.0


def _extra(artifact:Path)->tuple[float,list[str],float]:
    allowed={*REQUIRED,"solve.py","solve.sh","__pycache__"};extras=sorted(x.name for x in artifact.iterdir() if x.name not in allowed);return (1.0 if extras else 0.0),([f"存在未要求文件：{extras}"] if extras else []),1.0


def _safe_positive(call,*args):
    try:return call(*args)
    except Exception as exc:return 0.0,[f"任务专属检查异常：{type(exc).__name__}: {exc}"],0,1


def _safe_penalty(call,*args):
    try:return call(*args)
    except Exception as exc:return 1.0,[f"严重边界无法核验：{type(exc).__name__}: {exc}"],1.0


def evaluate(artifact:Path)->dict[str,Any]:
    ok,errors=_core(artifact);hashes=[_sha256(artifact/name) for name in REQUIRED if (artifact/name).is_file()]
    if not ok:return _fatal(errors,hashes)
    truth=V1.derive_truth(_input_dir());payload=json.loads((artifact/RESULT_JSON).read_text(encoding="utf-8"));results=[_criterion("R001",1,[],5,5)];scores={"R001":1.0}
    canonical=_safe_positive(_run_processor,artifact/PROCESSOR,_input_dir(),truth);mutation=_safe_positive(_mutation,artifact,_input_dir());json_check=_ratio(_atomic(payload,truth,"JSON"));excel=_read_excel_atomic(artifact,truth);word=_safe_positive(_word_atomic,artifact,truth);email=_safe_positive(_email_atomic,artifact,truth)
    word_score,word_failures,word_earned,word_total,word_text=word;email_score,email_failures,email_earned,email_total,email_text=email
    cross=_safe_positive(_cross,truth,payload,excel,word_text,email_text);format_check=_ratio([*excel["format"],(len(V1.extract_docx(artifact/WORD)[2])>=3,"Word 原生表格少于3张"),(V1.extract_docx(artifact/WORD)[3]==0,"Word 用图片替代内容")]);render=_safe_positive(_render,artifact)
    checks={"R002":canonical,"R003":mutation,"R004":json_check,"R005":excel["shipping"],"R006":excel["tax"],"R007":excel["formula"],"R008":word[:4],"R009":email[:4],"R010":cross,"R011":format_check,"R012":render}
    for key,(score,failures,earned,total) in checks.items():results.append(_criterion(key,score,failures,earned,total));scores[key]=score
    penalties={"P001":_safe_penalty(_unsupported,word_text,email_text,truth),"P002":_safe_penalty(_extra,artifact)}
    for key,(severity,labels,total) in penalties.items():results.append(_criterion(key,severity,labels,severity*total,total));scores[key]=severity
    total=sum(POSITIVE_WEIGHTS.values());positive=sum(POSITIVE_WEIGHTS[k]*scores[k] for k in POSITIVE_WEIGHTS);deduction=sum(abs(PENALTY_WEIGHTS[k])*scores[k] for k in PENALTY_WEIGHTS);reward=round(max(0.0,min(1.0,(positive-deduction)/total)),6)
    return {"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":hashes,"criterion_results":results,"reward":reward,"passed":reward>=0.8,"hurdle_failures":[],"penalties":[{"id":k,"amount":round(abs(PENALTY_WEIGHTS[k])*scores[k]/total,6)} for k in PENALTY_WEIGHTS if scores[k]>0]}


def main()->None:
    parser=argparse.ArgumentParser();parser.add_argument("--artifact-dir",type=Path,required=True);parser.add_argument("--result",type=Path,required=True);parser.add_argument("--reward-file",type=Path,required=True);parser.add_argument("--scoring-version",choices=("v1","v2"),default="v2");args=parser.parse_args()
    try:result=V1.evaluate(args.artifact_dir,render=True) if args.scoring_version=="v1" else evaluate(args.artifact_dir)
    except Exception as exc:result=_fatal([f"Judge 异常：{type(exc).__name__}: {exc}"],[])
    rendered=json.dumps(result,ensure_ascii=False,indent=2);args.result.parent.mkdir(parents=True,exist_ok=True);args.reward_file.parent.mkdir(parents=True,exist_ok=True);args.result.write_text(rendered+"\n",encoding="utf-8");args.reward_file.write_text(f"{float(result['reward']):.6f}\n",encoding="utf-8");print(rendered)


if __name__=="__main__":main()
