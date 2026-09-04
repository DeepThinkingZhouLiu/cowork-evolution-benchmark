#!/usr/bin/env python3
"""配送与税率更新传播任务的 v2 分层负例。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Callable

from docx import Document
from openpyxl import load_workbook
from openpyxl.styles import PatternFill


ROOT=Path(__file__).resolve().parents[1]
GOLDEN=ROOT/"solution"
INPUTS=ROOT/"environment"/"input_files"
REQUIRED=("shipping_tax_update_processor.py","shipping_tax_results.json","Shipping_Tax_Analysis.xlsx","Operations_Report.docx","operations_email.txt")
Mutator=Callable[[Path],None]


def _load_judge():
    path=Path(__file__).with_name("judge.py");spec=importlib.util.spec_from_file_location("shipping_tax_v2_judge_for_negatives",path)
    if spec is None or spec.loader is None:raise RuntimeError("无法载入 v2 Judge")
    module=importlib.util.module_from_spec(spec);sys.modules[spec.name]=module;spec.loader.exec_module(module);return module


JUDGE=_load_judge()


def _copy(target:Path)->None:
    target.mkdir(parents=True,exist_ok=True)
    for name in (*REQUIRED,"solve.py","solve.sh"):
        source=GOLDEN/name
        if source.is_file():shutil.copy2(source,target/name)


def _json(mutator:Callable[[dict],None])->Mutator:
    def apply(root:Path)->None:
        path=root/"shipping_tax_results.json";payload=json.loads(path.read_text(encoding="utf-8"));mutator(payload);path.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    return apply


def _excel(sheet:str,cell:str,value:object)->Mutator:
    def apply(root:Path)->None:
        path=root/"Shipping_Tax_Analysis.xlsx";book=load_workbook(path);book[sheet][cell]=value;book.save(path);book.close()
    return apply


def _processor_wrong_enabled(root:Path)->None:
    path=root/"shipping_tax_update_processor.py";text=path.read_text(encoding="utf-8");path.write_text(text.replace('enabled = [item for item in methods_by_zone[zone_id] if item["enabled"]]','enabled = list(methods_by_zone[zone_id])[:1]'),encoding="utf-8")


def _processor_hardcoded(root:Path)->None:
    script=(
        "#!/usr/bin/env python3\nimport argparse,shutil\nfrom pathlib import Path\n"
        "p=argparse.ArgumentParser();p.add_argument('--input-dir');p.add_argument('--output-dir',type=Path);a=p.parse_args();a.output_dir.mkdir(parents=True,exist_ok=True)\n"
        "source=Path(__file__).resolve().parent\n"
        "for name in ('shipping_tax_results.json','Shipping_Tax_Analysis.xlsx','Operations_Report.docx','operations_email.txt'):shutil.copy2(source/name,a.output_dir/name)\n"
    );(root/"shipping_tax_update_processor.py").write_text(script,encoding="utf-8")


def _word_heading(root:Path)->None:
    path=root/"Operations_Report.docx";doc=Document(path)
    for paragraph in doc.paragraphs:
        if paragraph.text.strip()=="Shipping Coverage":paragraph.text="Coverage";break
    doc.save(path)


def _word_table(root:Path,first:str,column:int,value:str)->None:
    path=root/"Operations_Report.docx";doc=Document(path)
    for table in doc.tables:
        for row in table.rows:
            if row.cells[0].text.strip()==first and len(row.cells)>column:row.cells[column].text=value;doc.save(path);return
    raise AssertionError(f"Word 未找到行 {first}")


def _word_boundary(root:Path)->None:
    path=root/"Operations_Report.docx";doc=Document(path)
    for paragraph in doc.paragraphs:
        if "No order volume, revenue, customer demand" in paragraph.text:paragraph.text="This report uses the supplied files.";break
    doc.save(path)


def _word_unsupported(root:Path)->None:
    path=root/"Operations_Report.docx";doc=Document(path);doc.add_paragraph("Revenue will increase by 12% after this configuration change.");doc.save(path)


def _email_replace(old:str,new:str)->Mutator:
    def apply(root:Path)->None:
        path=root/"operations_email.txt";text=path.read_text(encoding="utf-8");assert old in text;path.write_text(text.replace(old,new,1),encoding="utf-8")
    return apply


def _format_header(root:Path)->None:
    path=root/"Shipping_Tax_Analysis.xlsx";book=load_workbook(path);book["Shipping Zones"]["A1"].fill=PatternFill(fill_type=None);book.save(path);book.close()


def _moderate_json_conflict(root:Path)->None:
    path=root/"shipping_tax_results.json";value=json.loads(path.read_text(encoding="utf-8"))
    for zone in value["shipping_zones"]:zone["methods_count"]+=2
    for tax_class in value["tax_classes"]:
        for rate in tax_class["rates"]:rate["rate"]=float(rate["rate"])+1
    value["summary"].update({"total_enabled_methods":12,"configured_tax_rates":8,"highest_configured_rate":9.625})
    path.write_text(json.dumps(value,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")


def _severe_valid_incomplete(root:Path)->None:
    json_path=root/"shipping_tax_results.json";value=json.loads(json_path.read_text(encoding="utf-8"));value["shipping_zones"]=[];value["tax_classes"]=[];value["summary"]={};json_path.write_text(json.dumps(value,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    excel_path=root/"Shipping_Tax_Analysis.xlsx";book=load_workbook(excel_path)
    for sheet in book.worksheets:
        if sheet.max_row>1:sheet.delete_rows(2,sheet.max_row-1)
    book.save(excel_path);book.close()
    word_path=root/"Operations_Report.docx";doc=Document(word_path)
    for table in doc.tables:
        for row in table.rows[1:]:
            for cell in row.cells:cell.text=""
    for paragraph in doc.paragraphs:
        if paragraph.style.name=="Normal":paragraph.text="Source snapshot received; detailed analysis is incomplete."
    doc.save(word_path)
    (root/"operations_email.txt").write_text("To: operations@store.com\nSubject: Shipping and Tax Configuration Report\n\nAnalysis is incomplete; no configuration facts are supplied.\n",encoding="utf-8")


def _moderate_shipping_domain_incomplete(root:Path)->None:
    json_path=root/"shipping_tax_results.json";value=json.loads(json_path.read_text(encoding="utf-8"));value["shipping_zones"]=[];value["summary"]["total_zones"]=0;value["summary"]["total_enabled_methods"]=0;value["summary"]["zones_with_single_method"]=0;json_path.write_text(json.dumps(value,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    excel_path=root/"Shipping_Tax_Analysis.xlsx";book=load_workbook(excel_path);sheet=book["Shipping Zones"];sheet.delete_rows(2,sheet.max_row-1);book.save(excel_path);book.close()
    word_path=root/"Operations_Report.docx";doc=Document(word_path)
    shipping_table=doc.tables[1]
    for row in shipping_table.rows[1:]:
        for cell in row.cells:cell.text=""
    doc.save(word_path)
    email_path=root/"operations_email.txt";lines=email_path.read_text(encoding="utf-8").splitlines();lines=[line for line in lines if not any(zone in line for zone in ("Доставка по РФ","Москва","Международная"))];email_path.write_text("\n".join(lines)+"\n",encoding="utf-8")


def _input_only(root:Path)->None:
    for path in list(root.iterdir()):
        if path.is_file():path.unlink()
    for source in INPUTS.iterdir():
        if source.is_file():shutil.copy2(source,root/source.name)


PARTIALS={
    "json_one_zone_name_wrong":_json(lambda v:v["shipping_zones"][0].update({"zone_name":"Domestic US"})),
    "json_one_location_code_wrong":_json(lambda v:v["shipping_zones"][1]["locations"][0].update({"code":"US:NY"})),
    "json_one_method_setting_wrong":_json(lambda v:v["shipping_zones"][0]["enabled_methods"][0]["settings"].update({"cost":"10.99"})),
    "json_one_tax_rate_wrong":_json(lambda v:v["tax_classes"][0]["rates"][0].update({"rate":6})),
    "json_one_summary_wrong":_json(lambda v:v["summary"].update({"configured_tax_rates":8})),
    "processor_one_enabled_filter_wrong":_processor_wrong_enabled,
    "processor_hardcoded_against_rate_change":_processor_hardcoded,
    "excel_one_zone_name_wrong":_excel("Shipping Zones","A2","Domestic US"),
    "excel_one_region_wrong":_excel("Shipping Zones","C3","state:US:NY"),
    "excel_one_method_missing":_excel("Shipping Zones","D2","Стандартная доставка; Экспресс-доставка"),
    "excel_one_tax_rate_wrong":_excel("Tax Rates","C2",6),
    "excel_one_tax_state_wrong":_excel("Tax Rates","B4","NV"),
    "excel_zero_rate_not_blank":_excel("Tax Rates","C9",0),
    "excel_one_count_formula_flattened":_excel("Shipping Zones","B2",3),
    "excel_one_pseudo_dynamic_formula":_excel("Shipping Zones","B3",'=IF(D3="",0,2)'),
    "excel_one_header_fill_missing":_format_header,
    "word_one_heading_missing":_word_heading,
    "word_one_zone_method_count_wrong":lambda root:_word_table(root,"Москва",1,"3"),
    "word_one_tax_rate_wrong":lambda root:_word_table(root,"Reduced Rate",3,"6%"),
    "word_source_boundary_missing":_word_boundary,
    "word_unsupported_prediction":_word_unsupported,
    "email_wrong_recipient":_email_replace("operations@store.com","ops@store.com"),
    "email_wrong_subject":_email_replace("Shipping and Tax Configuration Report","Shipping Configuration Report"),
    "email_one_zone_count_wrong":_email_replace("Москва: 2 enabled","Москва: 3 enabled"),
    "email_one_tax_rate_wrong":_email_replace("rates 5%","rates 6%"),
    "moderate_json_cross_artifact_conflict":_moderate_json_conflict,
    "moderate_shipping_domain_incomplete":_moderate_shipping_domain_incomplete,
    "severe_valid_but_incomplete_package":_severe_valid_incomplete,
    "unexpected_debug_file":lambda root:(root/"debug-notes.txt").write_text("unrequested\n",encoding="utf-8"),
}


def main()->None:
    parser=argparse.ArgumentParser();parser.add_argument("--report",type=Path);args=parser.parse_args();render=JUDGE._render(GOLDEN);assert render[0]==1,render;assert JUDGE.evaluate(GOLDEN)["reward"]==1.0;JUDGE._render=lambda _artifact:render;records=[]
    with tempfile.TemporaryDirectory(prefix="shipping-tax-v2-negatives-") as raw:
        base=Path(raw)
        for name in REQUIRED:
            for mode in ("missing","empty","corrupt"):
                case=base/f"fatal-{mode}-{name}";_copy(case);target=case/name
                if mode=="missing":target.unlink()
                elif mode=="empty":target.write_bytes(b"")
                else:target.write_bytes(b"not-a-valid-task-artifact")
                reward=float(JUDGE.evaluate(case)["reward"]);assert reward==0,(mode,name,reward);records.append({"case":f"{mode}:{name}","layer":"fatal_zero","reward":reward})
        empty=base/"fatal-empty";empty.mkdir();reward=float(JUDGE.evaluate(empty)["reward"]);assert reward==0;records.append({"case":"empty-output","layer":"fatal_zero","reward":reward})
        input_only=base/"fatal-input-only";_copy(input_only);_input_only(input_only);reward=float(JUDGE.evaluate(input_only)["reward"]);assert reward==0;records.append({"case":"unmodified-input-only","layer":"fatal_zero","reward":reward})
        for name,mutator in PARTIALS.items():
            case=base/f"partial-{name}";_copy(case);mutator(case);reward=float(JUDGE.evaluate(case)["reward"]);assert 0<reward<1,(name,reward);records.append({"case":name,"layer":"partial","reward":reward})
    partial=[r["reward"] for r in records if r["layer"]=="partial"];report={"total":len(records),"fatal_zero":sum(r["layer"]=="fatal_zero" for r in records),"partial_continuous":len(partial),"partial_reward_range":[min(partial),max(partial)],"cases":records}
    if args.report:args.report.parent.mkdir(parents=True,exist_ok=True);args.report.write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=="__main__":main()
