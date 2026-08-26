#!/usr/bin/env python3
"""Golden、fresh replay 和七个必拒负例。"""

from __future__ import annotations

import importlib.util
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


TASK_ROOT = Path(__file__).resolve().parents[1]
SOLUTION = TASK_ROOT / "solution"
INPUTS = TASK_ROOT / "data" / "input_files"
OUTPUT_NAMES = [
    "shipping_tax_update_processor.py",
    "shipping_tax_results.json",
    "Shipping_Tax_Analysis.xlsx",
    "Operations_Report.docx",
    "operations_email.txt",
]

spec = importlib.util.spec_from_file_location("shipping_tax_judge", Path(__file__).with_name("judge.py"))
if spec is None or spec.loader is None:
    raise RuntimeError("无法加载 Judge")
judge = importlib.util.module_from_spec(spec)
spec.loader.exec_module(judge)


def copy_candidate(target: Path) -> None:
    target.mkdir(parents=True, exist_ok=True)
    for name in OUTPUT_NAMES:
        shutil.copy2(SOLUTION / name, target / name)


def evaluate_fast(path: Path) -> dict:
    return judge.evaluate(path.resolve(), render=False)


def assert_full_score(label: str, path: Path) -> None:
    result = evaluate_fast(path)
    assert result["passed"] and result["reward"] == 1.0, (
        label,
        result["reward"],
        result["hurdle_failures"],
    )


def assert_rejected(label: str, path: Path, failed_criterion: str) -> None:
    result = evaluate_fast(path)
    criteria = {item["criterion_id"]: item for item in result["criterion_results"]}
    assert not result["passed"] and result["reward"] < 0.8, (label, result["reward"])
    assert criteria[failed_criterion]["score"] == 0.0, (
        label,
        failed_criterion,
        criteria[failed_criterion],
    )


def case_golden(base: Path) -> None:
    candidate = base / "golden"
    copy_candidate(candidate)
    assert_full_score("golden", candidate)


def case_fresh_replay(base: Path) -> None:
    candidate = base / "fresh"
    candidate.mkdir()
    shutil.copy2(SOLUTION / "shipping_tax_update_processor.py", candidate)
    subprocess.run(
        [
            sys.executable,
            str(candidate / "shipping_tax_update_processor.py"),
            "--input-dir",
            str(INPUTS),
            "--output-dir",
            str(candidate),
        ],
        check=True,
        timeout=120,
    )
    assert_full_score("fresh replay", candidate)


def case_missing(base: Path) -> None:
    candidate = base / "missing"
    copy_candidate(candidate)
    (candidate / "Operations_Report.docx").unlink()
    assert_rejected("missing required file", candidate, "R001")


def case_wrong_json(base: Path) -> None:
    candidate = base / "wrong-json"
    copy_candidate(candidate)
    path = candidate / "shipping_tax_results.json"
    payload = json.loads(path.read_text(encoding="utf-8"))
    payload["summary"]["total_zones"] = 99
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    assert_rejected("wrong JSON fact", candidate, "R004")


def case_wrong_excel_tax(base: Path) -> None:
    candidate = base / "wrong-excel-tax"
    copy_candidate(candidate)
    path = candidate / "Shipping_Tax_Analysis.xlsx"
    workbook = load_workbook(path)
    sheet = workbook["Tax Rates"]
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 1).value == "Reduced Rate":
            sheet.cell(row, 3).value = 4
            break
    workbook.save(path)
    workbook.close()
    assert_rejected("wrong Excel tax rate", candidate, "R006")


def case_hardcoded_processor(base: Path) -> None:
    candidate = base / "hardcoded-processor"
    copy_candidate(candidate)
    script = candidate / "shipping_tax_update_processor.py"
    script.write_text(
        """#!/usr/bin/env python3
import argparse
import shutil
from pathlib import Path

parser = argparse.ArgumentParser()
parser.add_argument('--input-dir', required=True)
parser.add_argument('--output-dir', required=True, type=Path)
args = parser.parse_args()
args.output_dir.mkdir(parents=True, exist_ok=True)
source = Path(__file__).resolve().parent
for name in ['shipping_tax_results.json', 'Shipping_Tax_Analysis.xlsx', 'Operations_Report.docx', 'operations_email.txt']:
    shutil.copy2(source / name, args.output_dir / name)
""",
        encoding="utf-8",
    )
    assert_rejected("hardcoded processor", candidate, "R003")


def case_hardcoded_formula(base: Path) -> None:
    candidate = base / "hardcoded-formula"
    copy_candidate(candidate)
    path = candidate / "Shipping_Tax_Analysis.xlsx"
    workbook = load_workbook(path)
    sheet = workbook["Shipping Zones"]
    counts = [3, 2, 1]
    for row, count in enumerate(counts, start=2):
        sheet.cell(row, 2).value = f'=IF(D{row}="",0,{count})'
    workbook.save(path)
    workbook.close()
    assert_rejected("hardcoded pseudo-dynamic formula", candidate, "R007")


def case_word_conflict(base: Path) -> None:
    candidate = base / "word-conflict"
    copy_candidate(candidate)
    path = candidate / "Operations_Report.docx"
    document = Document(path)
    changed = False
    for table in document.tables:
        for row in table.rows:
            if row.cells[0].text.strip() == "Reduced Rate" and len(row.cells) >= 4:
                row.cells[3].text = "6%"
                changed = True
                break
        if changed:
            break
    assert changed
    document.save(path)
    assert_rejected("cross-file Word conflict", candidate, "R008")


def case_unsupported_assertion(base: Path) -> None:
    candidate = base / "unsupported"
    copy_candidate(candidate)
    path = candidate / "Operations_Report.docx"
    document = Document(path)
    document.add_paragraph("Revenue will increase by 12% after this change.")
    document.save(path)
    assert_rejected("unsupported business assertion", candidate, "R011")


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="cowork-evo-shipping-tax-regression-") as temp_dir:
        base = Path(temp_dir)
        case_golden(base)
        case_fresh_replay(base)
        case_missing(base)
        case_wrong_json(base)
        case_wrong_excel_tax(base)
        case_hardcoded_processor(base)
        case_hardcoded_formula(base)
        case_word_conflict(base)
        case_unsupported_assertion(base)
    print("PASS: Golden=1.0, fresh replay=1.0, 7/7 negative cases rejected")


if __name__ == "__main__":
    main()
