#!/usr/bin/env python3
"""生成学生风险分析的 Excel、Word、Tracker 和邮件草稿。"""

from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


SCRIPT_NAME = "risk_scorer.py"
JSON_NAME = "risk_analysis_results.json"
XLSX_NAME = "Student_Risk_Analysis.xlsx"
DOCX_NAME = "Intervention_Plan.docx"
TRACKER_NAME = "Student_Risk_Tracker.md"
EMAIL_NAME = "advisor_email.txt"
NAVY, BLUE, PALE, WHITE, INK, MUTED = "17365D", "2F75B5", "EAF2F8", "FFFFFF", "1F2937", "5B6B7F"


def add_table(sheet, ref: str, name: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)


def style_sheet(sheet, widths: dict[str, float]) -> None:
    fill = PatternFill("solid", fgColor=NAVY)
    thin = Side(style="thin", color="D7E0EA")
    for cell in sheet[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 28
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color=INK)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
        sheet.row_dimensions[row[0].row].height = 24
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.page_margins.left = 0.25; sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.35; sheet.page_margins.bottom = 0.35
    sheet.oddFooter.center.text = "Cowork-Bench · Student Retention"


def write_workbook(path: Path, result: dict) -> None:
    workbook = Workbook()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    overview = workbook.active
    overview.title = "Course_Overview"
    headers = ["course_id", "course_name", "enrollment_count", "avg_score", "pass_rate"]
    overview.append(headers)
    for row in result["course_overview"]:
        overview.append([row[header] for header in headers])
    style_sheet(overview, {"A": 13, "B": 39, "C": 20, "D": 15, "E": 15})
    for cell in overview["D"][1:]: cell.number_format = "0.00"
    for cell in overview["E"][1:]: cell.number_format = '0.00"%"'
    add_table(overview, "A1:E3", "Student_Course_Overview")

    distribution = workbook.create_sheet("Risk_Distribution")
    distribution.append(["risk_level", "student_count", "pct"])
    for index, row in enumerate(result["risk_distribution"], start=2):
        distribution.append([row["risk_level"], row["student_count"], None])
        distribution[f"C{index}"] = f"=ROUND(B{index}/SUM($B$2:$B$4)*100,2)"
    style_sheet(distribution, {"A": 18, "B": 20, "C": 16})
    for cell in distribution["C"][1:]: cell.number_format = '0.00"%"'
    add_table(distribution, "A1:C4", "Student_Risk_Distribution")

    at_risk = workbook.create_sheet("At_Risk_Students")
    headers = ["course_name", "high_risk_count", "medium_risk_count", "low_risk_count"]
    at_risk.append(headers)
    for row in result["at_risk_students"]:
        at_risk.append([row[header] for header in headers])
    style_sheet(at_risk, {"A": 39, "B": 21, "C": 23, "D": 20})
    add_table(at_risk, "A1:D3", "Student_At_Risk")

    intervention = workbook.create_sheet("Intervention_Plan")
    headers = ["risk_level", "action", "timeline", "responsible"]
    intervention.append(headers)
    for row in result["intervention_plan"]:
        intervention.append([row[header] for header in headers])
    style_sheet(intervention, {"A": 17, "B": 38, "C": 20, "D": 28})
    add_table(intervention, "A1:D4", "Student_Intervention_Plan")
    workbook.save(path)


def set_shading(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    props.append(shading)


def set_cell(cell, value, bold=False, color=INK, size=9) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(str(value))
    run.bold = bold; run.font.name = "Arial"; run.font.size = Pt(size); run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.7); section.bottom_margin = Cm(1.6); section.left_margin = Cm(1.8); section.right_margin = Cm(1.8)
    document.styles["Normal"].font.name = "Arial"; document.styles["Normal"].font.size = Pt(10); document.styles["Normal"].font.color.rgb = RGBColor.from_string(INK)
    document.styles["Normal"].paragraph_format.space_after = Pt(6)
    for name, size, color in (("Title", 26, NAVY), ("Heading 1", 17, NAVY), ("Heading 2", 12, BLUE)):
        style = document.styles[name]; style.font.name = "Arial"; style.font.size = Pt(size); style.font.bold = True; style.font.color.rgb = RGBColor.from_string(color)


def write_document(path: Path, result: dict) -> None:
    document = Document(); configure_document(document)
    document.add_paragraph("Student Retention Intervention Plan", style="Title")
    subtitle = document.add_paragraph("Foundations of Finance · Fall 2013 / Fall 2014")
    subtitle.runs[0].font.name = "Arial"; subtitle.runs[0].font.size = Pt(11); subtitle.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    distribution = {row["risk_level"]: row for row in result["risk_distribution"]}
    overview = result["course_overview"]
    total = sum(row["student_count"] for row in result["risk_distribution"])
    document.add_heading("Executive Summary", level=1)
    document.add_paragraph(
        f"Анализ охватывает два курса и {total:,} студентов хотя бы с одной оценённой работой. "
        f"Выявлено {distribution['High']['student_count']:,} High ({distribution['High']['pct']:.2f}%), "
        f"{distribution['Medium']['student_count']:,} Medium ({distribution['Medium']['pct']:.2f}%) и {distribution['Low']['student_count']:,} Low ({distribution['Low']['pct']:.2f}%). "
        "Приоритет — немедленно запустить индивидуальную работу с High-группой."
    )
    document.add_heading("Course Overview", level=1)
    table = document.add_table(rows=1, cols=6); table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Курс", "Enrollment", "Graded", "Submissions", "Avg", "Pass Rate"]
    for cell, value in zip(table.rows[0].cells, headers): set_shading(cell, NAVY); set_cell(cell, value, True, WHITE, 8)
    for row in overview:
        values = [row["course_name"], f"{row['enrollment_count']:,}", f"{row['graded_student_count']:,}", f"{row['graded_submission_count']:,}", f"{row['avg_score']:.2f}", f"{row['pass_rate']:.2f}%"]
        for cell, value in zip(table.add_row().cells, values): set_cell(cell, value, False, INK, 8)
    document.add_paragraph("Разница между Enrollment и Graded — это студенты без оценённых submission; они не включены в знаменатель риска и pass rate.")
    document.add_heading("Risk Distribution", level=1)
    table = document.add_table(rows=1, cols=3); table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, value in zip(table.rows[0].cells, ["Risk Level", "Student Count", "Share"]): set_shading(cell, NAVY); set_cell(cell, value, True, WHITE, 9)
    for row in result["risk_distribution"]:
        for cell, value in zip(table.add_row().cells, [row["risk_level"], f"{row['student_count']:,}", f"{row['pct']:.2f}%"]): set_cell(cell, value)
    document.add_page_break()
    document.add_heading("Recommended Interventions", level=1)
    counts = {row["risk_level"]: row["student_count"] for row in result["risk_distribution"]}
    for row in result["intervention_plan"]:
        document.add_heading(f"{row['risk_level']} · {counts[row['risk_level']]:,} студентов", level=2)
        document.add_paragraph(f"{row['action']}. Срок: {row['timeline']}. Ответственный: {row['responsible']}.")
    document.add_heading("Method & Scope", level=1)
    document.add_paragraph(
        "Для каждого course_id + user_id взято среднее арифметическое всех непустых score. High <60; Medium 60–74.99; Low ≥75. "
        "Course avg — среднее всех оценённых submission курса. Входы не содержат причин низкой успеваемости, результатов вмешательств или временного ряда; такие выводы не делались."
    )
    document.add_paragraph("Источники: canvas_courses.csv, canvas_enrollments.csv, canvas_assignments.csv, canvas_submissions.csv, scoring_model.json, Retention_Policy.pdf.")
    document.save(path)


def write_tracker(path: Path, result: dict) -> None:
    lines = ["# Student Risk Tracker", "", "| Student Course | Risk Level | Student Count | Average Score | Pass Rate |", "|---|---:|---:|---:|---:|"]
    for row in result["student_risk_tracker"]:
        lines.append(f"| {row['Student Course']} | {row['Risk Level']} | {row['Student Count']} | {row['Average Score']:.2f} | {row['Pass Rate']:.2f}% |")
    lines.extend(["", "数据口径：Student Count 仅包含至少有一条已评分 submission 的学生。"])
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_email(path: Path, result: dict) -> None:
    counts = {row["risk_level"]: row["student_count"] for row in result["risk_distribution"]}
    actions = {row["risk_level"]: row for row in result["intervention_plan"]}
    text = (
        "To: academic_advisors@university.edu\n"
        "Subject: Student Retention Risk Analysis - Action Required\n\n"
        "Коллеги,\n\n"
        f"Анализ Foundations of Finance Fall 2013/2014 выявил {counts['High']} студента High risk, {counts['Medium']} Medium и {counts['Low']} Low.\n\n"
        f"High: {actions['High']['action']}; {actions['High']['timeline']}; {actions['High']['responsible']}.\n"
        f"Medium: {actions['Medium']['action']}; {actions['Medium']['timeline']}; {actions['Medium']['responsible']}.\n"
        f"Low: {actions['Low']['action']}; {actions['Low']['timeline']}; {actions['Low']['responsible']}.\n\n"
        "Приложения подготовлены отдельно: Student_Risk_Analysis.xlsx и Intervention_Plan.docx.\n"
    )
    path.write_text(text, encoding="utf-8")


def solve(input_dir: Path, output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    source = Path(__file__).with_name(SCRIPT_NAME)
    script = output_dir / SCRIPT_NAME
    if source.resolve() != script.resolve(): shutil.copy2(source, script)
    subprocess.run([
        sys.executable, str(script),
        "--courses", str(input_dir / "canvas_courses.csv"), "--enrollments", str(input_dir / "canvas_enrollments.csv"),
        "--assignments", str(input_dir / "canvas_assignments.csv"), "--submissions", str(input_dir / "canvas_submissions.csv"),
        "--model", str(input_dir / "scoring_model.json"), "--output", str(output_dir / JSON_NAME),
    ], check=True, timeout=180)
    result = json.loads((output_dir / JSON_NAME).read_text(encoding="utf-8"))
    write_workbook(output_dir / XLSX_NAME, result)
    write_document(output_dir / DOCX_NAME, result)
    write_tracker(output_dir / TRACKER_NAME, result)
    write_email(output_dir / EMAIL_NAME, result)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=Path("/app/data/input_files"))
    parser.add_argument("--output-dir", type=Path, default=Path("/app/output"))
    args = parser.parse_args(); solve(args.input_dir, args.output_dir)


if __name__ == "__main__": main()
