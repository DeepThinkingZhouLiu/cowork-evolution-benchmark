#!/usr/bin/env python3
"""学生留存工作区任务的 v2 原子连续评分；v1 保存在 judge_v1.py。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import re
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from openpyxl import load_workbook


TASK_ID = "cowork-evo/cowork-bench-student-retention-workspace-package"
RUBRIC_VERSION = "cowork-bench-student-retention-workspace-package-v0.2"
JUDGE_VERSION = "student-retention-atomic-workspace-v2.0"
SCORER = "risk_scorer.py"
RESULT_JSON = "risk_analysis_results.json"
REPORT = "Student_Risk_Analysis.xlsx"
WORD = "Intervention_Plan.docx"
TRACKER = "Student_Risk_Tracker.md"
EMAIL = "advisor_email.txt"
REQUIRED = (SCORER, RESULT_JSON, REPORT, WORD, TRACKER, EMAIL)
POSITIVE_WEIGHTS = {
    "R001": 6, "R002": 11, "R003": 12, "R004": 8, "R005": 8,
    "R006": 10, "R007": 7, "R008": 10, "R009": 6, "R010": 6,
    "R011": 6, "R012": 4, "R013": 6,
}
PENALTY_WEIGHTS = {"P001": -12, "P002": -4}


def _load_v1():
    path = Path(__file__).with_name("judge_v1.py")
    spec = importlib.util.spec_from_file_location("student_retention_judge_v1", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v1 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


V1 = _load_v1()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _same(actual: Any, expected: Any) -> bool:
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        return V1.close(actual, expected)
    if expected is None:
        return actual in (None, "")
    return actual == expected


def _ratio(findings: Iterable[tuple[bool, str]]) -> tuple[float, list[str], int, int]:
    items = list(findings)
    earned = sum(bool(ok) for ok, _ in items)
    total = len(items)
    return earned / max(1, total), [label for ok, label in items if not ok], earned, total


def _criterion(identifier: str, score: float, failures: list[str], earned: float, total: float) -> dict[str, Any]:
    return {
        "criterion_id": identifier,
        "score": round(max(0.0, min(1.0, score)), 6),
        "evidence": "全部任务专属原子检查通过" if not failures else "未通过：" + "；".join(failures),
        "method": "deterministic",
        "raw": {"earned_units": round(earned, 6), "total_units": round(total, 6)},
    }


def _input_dir() -> Path:
    root = Path(__file__).resolve().parents[1]
    return Path("/app/input_files") if Path("/app/input_files").is_dir() else root / "data" / "input_files"


def _core(artifact: Path) -> tuple[bool, list[str]]:
    if not artifact.is_dir():
        return False, ["输出目录不存在"]
    errors: list[str] = []
    for name in REQUIRED:
        path = artifact / name
        if not path.is_file() or path.stat().st_size == 0:
            errors.append(f"核心交付缺失或为空：{name}")
            continue
        try:
            if name == SCORER:
                source = path.read_text(encoding="utf-8")
                compile(source, str(path), "exec")
                if len(source.strip()) < 100:
                    raise ValueError("评分器脚本过短")
            elif name == RESULT_JSON:
                if not isinstance(json.loads(path.read_text(encoding="utf-8")), dict):
                    raise ValueError("JSON 顶层不是对象")
            elif name == REPORT:
                if not zipfile.is_zipfile(path):
                    raise ValueError("不是有效 OOXML ZIP")
                book = load_workbook(path, read_only=True, data_only=False)
                if book.sheetnames != ["Course_Overview", "Risk_Distribution", "At_Risk_Students", "Intervention_Plan"]:
                    raise ValueError(f"工作表错误：{book.sheetnames}")
                book.close()
            elif name == WORD:
                if not zipfile.is_zipfile(path):
                    raise ValueError("不是有效 OOXML ZIP")
                document = Document(path)
                if not document.paragraphs or not document.tables:
                    raise ValueError("Word 无正文或原生表格")
            else:
                if len(path.read_text(encoding="utf-8").strip()) < 30:
                    raise ValueError("文本交付过短")
        except Exception as exc:
            errors.append(f"核心交付无法解析：{name}: {type(exc).__name__}: {exc}")
    return not errors, errors


def _fatal(errors: list[str], hashes: list[str]) -> dict[str, Any]:
    identifiers = [*POSITIVE_WEIGHTS, *PENALTY_WEIGHTS]
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes or ["0" * 64],
        "criterion_results": [_criterion(key, 0, errors, 0, 0) for key in identifiers],
        "reward": 0.0,
        "passed": False,
        "hurdle_failures": ["R001"],
        "penalties": [],
    }


def _atomic(actual: Any, expected: Any, path: str) -> list[tuple[bool, str]]:
    findings: list[tuple[bool, str]] = []
    if isinstance(expected, dict):
        source = actual if isinstance(actual, dict) else {}
        findings.extend([
            (isinstance(actual, dict), f"{path} 应为对象"),
            (list(source) == list(expected), f"{path} 键或顺序错误"),
        ])
        for key, value in expected.items():
            findings.extend(_atomic(source.get(key), value, f"{path}.{key}"))
    elif isinstance(expected, list):
        rows = actual if isinstance(actual, list) else []
        findings.extend([
            (isinstance(actual, list), f"{path} 应为数组"),
            (len(rows) == len(expected), f"{path} 数量错误"),
        ])
        for index, value in enumerate(expected):
            found = rows[index] if index < len(rows) else None
            findings.extend(_atomic(found, value, f"{path}[{index}]"))
    else:
        findings.append((_same(actual, expected), f"{path} 错误：实际={actual!r}，期望={expected!r}"))
    return findings


def _canonical(script: Path, input_dir: Path, truth: dict[str, Any]) -> tuple[float, list[str], int, int]:
    payload, log = V1.run_scorer(script, input_dir)
    findings = [(payload is not None, f"评分器回放失败：{log}")]
    findings.extend(_atomic(payload, truth, "scorer.JSON"))
    return _ratio(findings)


def _mutation(script: Path, input_dir: Path, truth_with_details: dict[str, Any]) -> tuple[float, list[str], int, int]:
    with tempfile.TemporaryDirectory(prefix="retention-v2-mutation-") as raw:
        mutated = Path(raw) / "inputs"
        course_id, user_id, changed = V1.build_mutated_inputs(input_dir, mutated, truth_with_details)
        expected = V1.derive_truth(mutated)
        actual, log = V1.run_scorer(script, mutated)
    baseline = V1.public_truth(truth_with_details)
    findings = [
        (actual is not None, f"变异回放失败：{log}"),
        (course_id == 17, f"变异课程错误：{course_id}"),
        (user_id == 5906, f"变异学生错误：{user_id}"),
        (changed == 11, f"变异非空分数行数错误：{changed}"),
        (expected["course_overview"] != baseline["course_overview"], "变异后课程概览未变化"),
        (expected["risk_distribution"] != baseline["risk_distribution"], "变异后总体风险未变化"),
        (expected["at_risk_students"] != baseline["at_risk_students"], "变异后分课程风险未变化"),
    ]
    findings.extend(_atomic(actual, expected, "mutated_scorer.JSON"))
    return _ratio(findings)


def _excel_atomic(artifact: Path, truth: dict[str, Any]) -> dict[str, Any]:
    path = artifact / REPORT
    book = load_workbook(path, data_only=False)
    specs = {
        "Course_Overview": (
            V1.OVERVIEW_HEADERS,
            [{key: row[key] for key in V1.OVERVIEW_HEADERS} for row in truth["course_overview"]],
        ),
        "Risk_Distribution": (
            V1.DISTRIBUTION_HEADERS,
            [{"risk_level": row["risk_level"], "student_count": row["student_count"]} for row in truth["risk_distribution"]],
        ),
        "At_Risk_Students": (V1.AT_RISK_HEADERS, truth["at_risk_students"]),
        "Intervention_Plan": (V1.INTERVENTION_HEADERS, truth["intervention_plan"]),
    }
    content_findings: list[tuple[bool, str]] = []
    actual_by_sheet: dict[str, list[dict[str, Any]]] = {}
    formatting: list[tuple[bool, str]] = []
    for sheet_name, (headers, expected_rows) in specs.items():
        sheet = book[sheet_name]
        actual_headers = [sheet.cell(1, index).value for index in range(1, len(headers) + 1)]
        content_findings.extend((actual == expected, f"{sheet_name} 表头第{index}列错误") for index, (actual, expected) in enumerate(zip(actual_headers, headers), 1))
        actual_rows: list[dict[str, Any]] = []
        for row_number in range(2, sheet.max_row + 1):
            values = [sheet.cell(row_number, index).value for index in range(1, len(headers) + 1)]
            if all(value in (None, "") for value in values):
                continue
            if sheet_name == "Risk_Distribution":
                actual_rows.append({"risk_level": values[0], "student_count": values[1]})
            else:
                actual_rows.append(dict(zip(headers, values)))
        actual_by_sheet[sheet_name] = actual_rows
        content_findings.extend(_atomic(actual_rows, expected_rows, f"Excel.{sheet_name}"))
        header_cells = [sheet.cell(1, index) for index in range(1, len(headers) + 1)]
        formatting.extend([
            (len(sheet.tables) >= 1 or bool(sheet.auto_filter.ref), f"{sheet_name} 无原生表或筛选"),
            (str(sheet.freeze_panes or "") == "A2", f"{sheet_name} 未冻结首行"),
            (all(bool(cell.font.bold) for cell in header_cells), f"{sheet_name} 表头未全部加粗"),
            (all(cell.fill.fill_type == "solid" for cell in header_cells), f"{sheet_name} 表头未全部填色"),
            (all((sheet.column_dimensions[cell.column_letter].width or 0) >= 10 for cell in header_cells), f"{sheet_name} 存在过窄列"),
            (len(sheet._images) == 0, f"{sheet_name} 用图片替代数据"),
        ])
    formulas = [book["Risk_Distribution"][cell].value for cell in ("C2", "C3", "C4")]
    formula_findings: list[tuple[bool, str]] = []
    for row, formula in enumerate(formulas, 2):
        formula_findings.extend([
            (isinstance(formula, str) and formula.startswith("="), f"C{row} 不是公式"),
            (isinstance(formula, str) and bool(re.search(rf"\bB{row}\b", formula, re.I)), f"C{row} 未引用 B{row}"),
            (isinstance(formula, str) and "SUM" in formula.upper() and "$B$2:$B$4" in formula.upper(), f"C{row} 未引用风险人数总和"),
        ])
    with tempfile.TemporaryDirectory(prefix="retention-v2-recalc-") as raw:
        recalculated, log = V1.recalculate_excel(path, Path(raw))
    expected_pct = [row["pct"] for row in truth["risk_distribution"]]
    formula_findings.append((recalculated is not None, f"LibreOffice 重算失败：{log}"))
    for index, expected in enumerate(expected_pct):
        actual = recalculated[index] if recalculated is not None and index < len(recalculated) else None
        formula_findings.append((_same(actual, expected), f"C{index + 2} 重算={actual!r}，期望={expected}"))
    book.close()
    return {
        "content": _ratio(content_findings),
        "formula": _ratio(formula_findings),
        "format": formatting,
        "rows": actual_by_sheet,
    }


def _cell_matches(value: str, expected: Any, percent: bool = False) -> bool:
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        cleaned = value.replace(",", "").replace("%", "").strip()
        return _same(cleaned, expected)
    return value.strip() == str(expected)


def _word_atomic(artifact: Path, truth: dict[str, Any]) -> tuple[float, list[str], int, int, str, int, int, list[str]]:
    path = artifact / WORD
    text, headings, table_count, media = V1.docx_payload(path)
    document = Document(path)
    tables = [[[cell.text.strip() for cell in row.cells] for row in table.rows] for table in document.tables]
    findings: list[tuple[bool, str]] = []
    heading_norm = {heading.casefold() for heading in headings}
    for heading in ("Student Retention Intervention Plan", "Executive Summary", "Course Overview", "Risk Distribution", "Recommended Interventions", "Method & Scope"):
        findings.append((heading.casefold() in heading_norm, f"Word 缺标题 {heading}"))
    overview_rows = tables[0][1:] if len(tables) >= 1 else []
    for expected in truth["course_overview"]:
        row = next((item for item in overview_rows if item and item[0] == expected["course_name"]), [])
        findings.append((bool(row), f"Word 缺课程 {expected['course_id']} 行"))
        wanted = [
            expected["course_name"], expected["enrollment_count"], expected["graded_student_count"],
            expected["graded_submission_count"], expected["avg_score"], expected["pass_rate"],
        ]
        for index, value in enumerate(wanted):
            actual = row[index] if index < len(row) else ""
            findings.append((_cell_matches(actual, value, percent=index == 5), f"Word 课程 {expected['course_id']} 第{index + 1}列错误"))
    risk_rows = tables[1][1:] if len(tables) >= 2 else []
    for expected in truth["risk_distribution"]:
        row = next((item for item in risk_rows if item and item[0] == expected["risk_level"]), [])
        findings.extend([
            (bool(row), f"Word 缺风险档 {expected['risk_level']}"),
            (len(row) > 1 and _cell_matches(row[1], expected["student_count"]), f"Word {expected['risk_level']} 人数错误"),
            (len(row) > 2 and _cell_matches(row[2], expected["pct"], percent=True), f"Word {expected['risk_level']} 百分比错误"),
        ])
    for expected in truth["intervention_plan"]:
        findings.extend([
            (expected["action"] in text, f"Word 缺 {expected['risk_level']} action"),
            (expected["timeline"] in text, f"Word 缺 {expected['risk_level']} timeline"),
            (expected["responsible"] in text, f"Word 缺 {expected['risk_level']} responsible"),
        ])
    for token, label in (("High <60", "High 阈值"), ("Medium 60", "Medium 阈值"), ("Low ≥75", "Low 阈值")):
        findings.append((token in text, f"Word 缺 {label}"))
    findings.extend([
        (any(token in text.casefold() for token in ("не содержат причин", "does not contain causes", "no causal")), "Word 缺无因果边界"),
        ("canvas_courses.csv" in text and "canvas_submissions.csv" in text and "scoring_model.json" in text, "Word 缺来源文件边界"),
    ])
    score, failures, earned, total = _ratio(findings)
    return score, failures, earned, total, text, table_count, media, headings


def _tracker_atomic(artifact: Path, truth: dict[str, Any]) -> tuple[float, list[str], int, int, str, list[list[str]]]:
    text = (artifact / TRACKER).read_text(encoding="utf-8")
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    table_lines = [line for line in lines if line.startswith("|") and line.endswith("|")]
    parsed = [[cell.strip() for cell in line.strip("|").split("|")] for line in table_lines]
    findings: list[tuple[bool, str]] = [
        (bool(lines) and lines[0] == "# Student Risk Tracker", "Markdown H1 错误"),
        (len(parsed) >= 4, "Markdown 表格不足两行"),
    ]
    expected_header = ["Student Course", "Risk Level", "Student Count", "Average Score", "Pass Rate"]
    actual_header = parsed[0] if parsed else []
    for index, expected in enumerate(expected_header):
        actual = actual_header[index] if index < len(actual_header) else ""
        findings.append((actual == expected, f"Markdown 表头第{index + 1}列错误"))
    data_rows = parsed[2:] if len(parsed) >= 2 else []
    findings.append((len(data_rows) == 2, f"Markdown 数据行数错误：{len(data_rows)}"))
    for row_index, expected in enumerate(truth["student_risk_tracker"]):
        row = data_rows[row_index] if row_index < len(data_rows) else []
        wanted = [
            expected["Student Course"], expected["Risk Level"], expected["Student Count"],
            expected["Average Score"], expected["Pass Rate"],
        ]
        for column, value in enumerate(wanted):
            actual = row[column] if column < len(row) else ""
            findings.append((_cell_matches(actual, value, percent=column == 4), f"Markdown 第{row_index + 1}行第{column + 1}列错误"))
    score, failures, earned, total = _ratio(findings)
    return score, failures, earned, total, text, data_rows


def _email_atomic(artifact: Path, truth: dict[str, Any]) -> tuple[float, list[str], int, int, str]:
    text = (artifact / EMAIL).read_text(encoding="utf-8")
    high_count = next(row["student_count"] for row in truth["risk_distribution"] if row["risk_level"] == "High")
    findings: list[tuple[bool, str]] = [
        (bool(re.search(r"^To:\s*academic_advisors@university\.edu\s*$", text, re.I | re.M)), "邮件收件人错误"),
        (bool(re.search(r"^Subject:\s*Student Retention Risk Analysis - Action Required\s*$", text, re.I | re.M)), "邮件主题错误"),
        (V1.contains_number(text, high_count), f"邮件 High 总人数错误：期望 {high_count}"),
    ]
    for expected in truth["intervention_plan"]:
        line = next((line for line in text.splitlines() if line.strip().startswith(expected["risk_level"] + ":")), "")
        findings.extend([
            (bool(line), f"邮件缺 {expected['risk_level']} 行"),
            (expected["action"] in line, f"邮件缺 {expected['risk_level']} action"),
            (expected["timeline"] in line, f"邮件缺 {expected['risk_level']} timeline"),
            (expected["responsible"] in line, f"邮件缺 {expected['risk_level']} responsible"),
        ])
    findings.extend([
        (REPORT in text, "邮件未提及 Excel 附件"),
        (WORD in text, "邮件未提及 Word 附件"),
    ])
    score, failures, earned, total = _ratio(findings)
    return score, failures, earned, total, text


def _cross_atomic(
    truth: dict[str, Any], payload: dict[str, Any], excel_rows: dict[str, list[dict[str, Any]]],
    word_text: str, tracker_rows: list[list[str]], email_text: str,
) -> tuple[float, list[str], int, int]:
    findings: list[tuple[bool, str]] = []
    for index, expected in enumerate(truth["course_overview"]):
        json_row = payload.get("course_overview", [])[index] if index < len(payload.get("course_overview", [])) else {}
        excel_row = excel_rows.get("Course_Overview", [])[index] if index < len(excel_rows.get("Course_Overview", [])) else {}
        tracker = tracker_rows[index] if index < len(tracker_rows) else []
        findings.extend([
            (_same(json_row.get("avg_score"), expected["avg_score"]) and _same(excel_row.get("avg_score"), expected["avg_score"]), f"课程 {expected['course_id']} JSON/Excel 均分冲突"),
            (len(tracker) > 3 and _cell_matches(tracker[3], expected["avg_score"]), f"课程 {expected['course_id']} Tracker 均分冲突"),
            (expected["course_name"] in word_text and V1.contains_number(word_text, expected["avg_score"]), f"课程 {expected['course_id']} Word 均分冲突"),
        ])
    for index, expected in enumerate(truth["risk_distribution"]):
        json_row = payload.get("risk_distribution", [])[index] if index < len(payload.get("risk_distribution", [])) else {}
        excel_row = excel_rows.get("Risk_Distribution", [])[index] if index < len(excel_rows.get("Risk_Distribution", [])) else {}
        count = expected["student_count"]
        findings.extend([
            (_same(json_row.get("student_count"), count) and _same(excel_row.get("student_count"), count), f"{expected['risk_level']} JSON/Excel 人数冲突"),
            (expected["risk_level"] in word_text and V1.contains_number(word_text, count), f"{expected['risk_level']} Word 人数冲突"),
            (expected["risk_level"] in email_text and V1.contains_number(email_text, count), f"{expected['risk_level']} 邮件人数冲突"),
        ])
    for index, expected in enumerate(truth["intervention_plan"]):
        json_row = payload.get("intervention_plan", [])[index] if index < len(payload.get("intervention_plan", [])) else {}
        excel_row = excel_rows.get("Intervention_Plan", [])[index] if index < len(excel_rows.get("Intervention_Plan", [])) else {}
        findings.extend([
            (json_row.get("action") == expected["action"] and excel_row.get("action") == expected["action"], f"{expected['risk_level']} JSON/Excel action 冲突"),
            (expected["action"] in word_text and expected["timeline"] in word_text, f"{expected['risk_level']} Word 干预冲突"),
            (expected["action"] in email_text and expected["timeline"] in email_text, f"{expected['risk_level']} 邮件干预冲突"),
        ])
    return _ratio(findings)


def _render(artifact: Path) -> tuple[float, list[str], int, int]:
    with tempfile.TemporaryDirectory(prefix="retention-v2-render-") as raw:
        rendered, log = V1.render_office([artifact / REPORT, artifact / WORD], Path(raw))
    excel = rendered.get(REPORT, {})
    word = rendered.get(WORD, {})
    return _ratio([
        (excel.get("pages", 0) >= 4, f"Excel 渲染少于4页：{log}"),
        (word.get("pages", 0) >= 2, f"Word 渲染少于2页：{log}"),
        (all(token in excel.get("text", "") for token in ("course_id", "risk_level", "high_risk_count", "responsible")), "Excel 渲染缺核心文字"),
        (all(token in word.get("text", "") for token in ("Executive Summary", "Risk Distribution", "Method & Scope")), "Word 渲染缺核心章节"),
    ])


def _unsupported(word_text: str, email_text: str, truth: dict[str, Any]) -> tuple[float, list[str], float]:
    labels = V1.unsupported_claims(word_text + "\n" + email_text, truth)
    severity = min(1.0, 0.45 + 0.15 * len(labels)) if labels else 0.0
    return severity, labels, 1.0


def _extra(artifact: Path) -> tuple[float, list[str], float]:
    allowed = {*REQUIRED, "solve.py", "solve.sh", "__pycache__"}
    extras = sorted(item.name for item in artifact.iterdir() if item.name not in allowed)
    return (1.0 if extras else 0.0), ([f"存在未要求文件：{extras}"] if extras else []), 1.0


def _safe_positive(call, *args):
    try:
        return call(*args)
    except Exception as exc:
        return 0.0, [f"任务专属检查异常：{type(exc).__name__}: {exc}"], 0, 1


def _safe_penalty(call, *args):
    try:
        return call(*args)
    except Exception as exc:
        return 1.0, [f"严重边界无法核验：{type(exc).__name__}: {exc}"], 1.0


def evaluate(artifact: Path) -> dict[str, Any]:
    ok, errors = _core(artifact)
    hashes = [_sha256(artifact / name) for name in REQUIRED if (artifact / name).is_file()]
    if not ok:
        return _fatal(errors, hashes)
    input_dir = _input_dir()
    truth_with_details = V1.derive_truth(input_dir, include_student_details=True)
    truth = V1.public_truth(truth_with_details)
    payload = json.loads((artifact / RESULT_JSON).read_text(encoding="utf-8"))
    results = [_criterion("R001", 1, [], len(REQUIRED), len(REQUIRED))]
    scores = {"R001": 1.0}

    canonical = _safe_positive(_canonical, artifact / SCORER, input_dir, truth)
    mutation = _safe_positive(_mutation, artifact / SCORER, input_dir, truth_with_details)
    overview = _ratio(_atomic(payload.get("course_overview"), truth["course_overview"], "JSON.course_overview"))
    remainder_expected = {key: truth[key] for key in ("risk_distribution", "at_risk_students", "intervention_plan", "student_risk_tracker")}
    remainder_actual = {key: payload.get(key) for key in remainder_expected}
    remainder = _ratio(_atomic(remainder_actual, remainder_expected, "JSON.remaining"))
    excel = _excel_atomic(artifact, truth)
    word = _safe_positive(_word_atomic, artifact, truth)
    tracker = _safe_positive(_tracker_atomic, artifact, truth)
    email = _safe_positive(_email_atomic, artifact, truth)
    word_score, word_failures, word_earned, word_total, word_text, word_tables, word_media, word_headings = word
    tracker_score, tracker_failures, tracker_earned, tracker_total, _tracker_text, tracker_rows = tracker
    email_score, email_failures, email_earned, email_total, email_text = email
    cross = _safe_positive(_cross_atomic, truth, payload, excel["rows"], word_text, tracker_rows, email_text)
    native = _ratio([
        *excel["format"],
        (word_tables >= 2, "Word 原生表格少于2张"),
        (word_media == 0, "Word 用图片替代内容"),
        (len(word_headings) >= 6, "Word 原生标题不足6个"),
    ])
    render = _safe_positive(_render, artifact)
    checks = {
        "R002": canonical,
        "R003": mutation,
        "R004": overview,
        "R005": remainder,
        "R006": excel["content"],
        "R007": excel["formula"],
        "R008": word[:4],
        "R009": tracker[:4],
        "R010": email[:4],
        "R011": cross,
        "R012": native,
        "R013": render,
    }
    for key, (score, failures, earned, total) in checks.items():
        results.append(_criterion(key, score, failures, earned, total))
        scores[key] = score
    penalties = {
        "P001": _safe_penalty(_unsupported, word_text, email_text, truth),
        "P002": _safe_penalty(_extra, artifact),
    }
    for key, (severity, labels, total) in penalties.items():
        results.append(_criterion(key, severity, labels, severity * total, total))
        scores[key] = severity
    positive_total = sum(POSITIVE_WEIGHTS.values())
    positive = sum(POSITIVE_WEIGHTS[key] * scores[key] for key in POSITIVE_WEIGHTS)
    deduction = sum(abs(PENALTY_WEIGHTS[key]) * scores[key] for key in PENALTY_WEIGHTS)
    reward = round(max(0.0, min(1.0, (positive - deduction) / positive_total)), 6)
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes,
        "criterion_results": results,
        "reward": reward,
        "passed": reward >= 0.8,
        "hurdle_failures": [],
        "penalties": [
            {"id": key, "amount": round(abs(PENALTY_WEIGHTS[key]) * scores[key] / positive_total, 6)}
            for key in PENALTY_WEIGHTS if scores[key] > 0
        ],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    parser.add_argument("--scoring-version", choices=("v1", "v2"), default="v2")
    args = parser.parse_args()
    try:
        result = V1.evaluate(args.artifact_dir) if args.scoring_version == "v1" else evaluate(args.artifact_dir)
    except Exception as exc:
        result = _fatal([f"Judge 异常：{type(exc).__name__}: {exc}"], [])
    rendered = json.dumps(result, ensure_ascii=False, indent=2)
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(rendered + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{float(result['reward']):.6f}\n", encoding="utf-8")
    print(rendered)


if __name__ == "__main__":
    main()
