#!/usr/bin/env python3
"""学生留存工作区任务的 v2 分层负例。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Callable

from docx import Document
from openpyxl import load_workbook
from openpyxl.styles import PatternFill


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution"
INPUTS = ROOT / "environment" / "input_files"
REQUIRED = (
    "risk_scorer.py", "risk_analysis_results.json", "Student_Risk_Analysis.xlsx",
    "Intervention_Plan.docx", "Student_Risk_Tracker.md", "advisor_email.txt",
)
Mutator = Callable[[Path], None]


def _load_judge():
    path = Path(__file__).with_name("judge.py")
    spec = importlib.util.spec_from_file_location("student_retention_v2_judge_for_negatives", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v2 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


JUDGE = _load_judge()


def _copy(target: Path) -> None:
    target.mkdir(parents=True, exist_ok=True)
    for name in (*REQUIRED, "solve.py", "solve.sh"):
        source = GOLDEN / name
        if source.is_file():
            shutil.copy2(source, target / name)


def _json(mutator: Callable[[dict], None]) -> Mutator:
    def apply(root: Path) -> None:
        path = root / "risk_analysis_results.json"
        payload = json.loads(path.read_text(encoding="utf-8"))
        mutator(payload)
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return apply


def _excel(sheet: str, cell: str, value: object) -> Mutator:
    def apply(root: Path) -> None:
        path = root / "Student_Risk_Analysis.xlsx"
        book = load_workbook(path)
        book[sheet][cell] = value
        book.save(path)
        book.close()
    return apply


def _scorer_threshold(root: Path) -> None:
    path = root / "risk_scorer.py"
    text = path.read_text(encoding="utf-8")
    old = "if value < high_cutoff:\n        return \"High\""
    assert old in text
    path.write_text(text.replace(old, "if value < high_cutoff + 5:\n        return \"High\"", 1), encoding="utf-8")


def _scorer_hardcoded(root: Path) -> None:
    script = (
        "#!/usr/bin/env python3\nimport argparse,shutil\nfrom pathlib import Path\n"
        "p=argparse.ArgumentParser();p.add_argument('--courses');p.add_argument('--enrollments');"
        "p.add_argument('--assignments');p.add_argument('--submissions');p.add_argument('--model');"
        "p.add_argument('--output',type=Path);a=p.parse_args();"
        "shutil.copy2(Path(__file__).with_name('risk_analysis_results.json'),a.output)\n"
    )
    (root / "risk_scorer.py").write_text(script, encoding="utf-8")


def _word_heading(root: Path) -> None:
    path = root / "Intervention_Plan.docx"
    document = Document(path)
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == "Risk Distribution":
            paragraph.text = "Risk Summary"
            break
    document.save(path)


def _word_table(root: Path, table_index: int, first: str, column: int, value: str) -> None:
    path = root / "Intervention_Plan.docx"
    document = Document(path)
    for row in document.tables[table_index].rows:
        if row.cells[0].text.strip() == first:
            row.cells[column].text = value
            document.save(path)
            return
    raise AssertionError(f"Word 未找到行 {first}")


def _word_replace(root: Path, old: str, new: str) -> None:
    path = root / "Intervention_Plan.docx"
    document = Document(path)
    for paragraph in document.paragraphs:
        if old in paragraph.text:
            paragraph.text = paragraph.text.replace(old, new, 1)
            document.save(path)
            return
    raise AssertionError(f"Word 未找到文本 {old}")


def _word_unsupported(root: Path) -> None:
    path = root / "Intervention_Plan.docx"
    document = Document(path)
    document.add_paragraph("Academic advising improved retention by 12% in the prior year.")
    document.save(path)


def _tracker_replace(old: str, new: str) -> Mutator:
    def apply(root: Path) -> None:
        path = root / "Student_Risk_Tracker.md"
        text = path.read_text(encoding="utf-8")
        assert old in text
        path.write_text(text.replace(old, new, 1), encoding="utf-8")
    return apply


def _email_replace(old: str, new: str) -> Mutator:
    def apply(root: Path) -> None:
        path = root / "advisor_email.txt"
        text = path.read_text(encoding="utf-8")
        assert old in text
        path.write_text(text.replace(old, new, 1), encoding="utf-8")
    return apply


def _format_header(root: Path) -> None:
    path = root / "Student_Risk_Analysis.xlsx"
    book = load_workbook(path)
    book["Course_Overview"]["A1"].fill = PatternFill(fill_type=None)
    book.save(path)
    book.close()


def _moderate_risk_domain_conflict(root: Path) -> None:
    json_path = root / "risk_analysis_results.json"
    payload = json.loads(json_path.read_text(encoding="utf-8"))
    wrong = [("High", 400, 10.77), ("Medium", 1100, 29.62), ("Low", 2214, 59.61)]
    payload["risk_distribution"] = [
        {"risk_level": level, "student_count": count, "pct": pct}
        for level, count, pct in wrong
    ]
    payload["at_risk_students"] = [
        {"course_name": "Основы финансов (Осень 2013)", "high_risk_count": 219, "medium_risk_count": 585, "low_risk_count": 1068},
        {"course_name": "Основы финансов (Осень 2014)", "high_risk_count": 181, "medium_risk_count": 515, "low_risk_count": 1146},
    ]
    for row in payload["course_overview"]:
        row["pass_rate"] = round(float(row["pass_rate"]) + 1, 2)
    for row in payload["student_risk_tracker"]:
        row["Risk Level"] = "Medium"
        row["Pass Rate"] = round(float(row["Pass Rate"]) + 1, 2)
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    excel_path = root / "Student_Risk_Analysis.xlsx"
    book = load_workbook(excel_path)
    sheet = book["Risk_Distribution"]
    for row, (_, count, _) in enumerate(wrong, 2):
        sheet[f"B{row}"] = count
    at_risk = book["At_Risk_Students"]
    for row, values in enumerate(((219, 585, 1068), (181, 515, 1146)), 2):
        for column, value in enumerate(values, 2):
            at_risk.cell(row, column).value = value
    book["Course_Overview"]["E2"] = 89.84
    book["Course_Overview"]["E3"] = 92.64
    book.save(excel_path)
    book.close()
    word_path = root / "Intervention_Plan.docx"
    document = Document(word_path)
    for row, (_, count, pct) in zip(document.tables[1].rows[1:], wrong):
        row.cells[1].text = f"{count:,}"
        row.cells[2].text = f"{pct:.2f}%"
    document.tables[0].rows[1].cells[5].text = "89.84%"
    document.tables[0].rows[2].cells[5].text = "92.64%"
    for paragraph in document.paragraphs:
        paragraph.text = paragraph.text.replace("363 High (9.77%), 1,161 Medium (31.26%) и 2,190 Low (58.97%)", "400 High (10.77%), 1,100 Medium (29.62%) и 2,214 Low (59.61%)")
    document.save(word_path)
    email_path = root / "advisor_email.txt"
    text = email_path.read_text(encoding="utf-8")
    text = text.replace("363 студента High risk, 1161 Medium и 2190 Low", "400 студентов High risk, 1100 Medium и 2214 Low")
    email_path.write_text(text, encoding="utf-8")
    tracker_path = root / "Student_Risk_Tracker.md"
    tracker_text = tracker_path.read_text(encoding="utf-8").replace(" | Low | ", " | Medium | ")
    tracker_text = tracker_text.replace("88.84%", "89.84%").replace("91.64%", "92.64%")
    tracker_path.write_text(tracker_text, encoding="utf-8")


def _moderate_risk_domain_incomplete(root: Path) -> None:
    json_path = root / "risk_analysis_results.json"
    payload = json.loads(json_path.read_text(encoding="utf-8"))
    payload["risk_distribution"] = []
    payload["at_risk_students"] = []
    payload["intervention_plan"] = []
    payload["student_risk_tracker"] = []
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    excel_path = root / "Student_Risk_Analysis.xlsx"
    book = load_workbook(excel_path)
    for name in ("Risk_Distribution", "At_Risk_Students", "Intervention_Plan"):
        sheet = book[name]
        if sheet.max_row > 1:
            sheet.delete_rows(2, sheet.max_row - 1)
    book.save(excel_path)
    book.close()
    word_path = root / "Intervention_Plan.docx"
    document = Document(word_path)
    for table in document.tables[1:2]:
        for row in table.rows[1:]:
            for cell in row.cells:
                cell.text = ""
    for paragraph in document.paragraphs:
        if paragraph.style.name == "Normal" and any(token in paragraph.text for token in ("High", "Medium", "Low", "One-on-one", "Group tutoring", "Self-paced", "<60", "≥75")):
            paragraph.text = "Risk analysis details are incomplete."
    document.save(word_path)
    (root / "Student_Risk_Tracker.md").write_text("# Student Risk Tracker\n\nThe risk-level rows are incomplete.\n", encoding="utf-8")
    (root / "advisor_email.txt").write_text(
        "To: academic_advisors@university.edu\nSubject: Student Retention Risk Analysis - Action Required\n\nRisk-level findings and interventions are incomplete.\n",
        encoding="utf-8",
    )


def _severe_valid_incomplete(root: Path) -> None:
    json_path = root / "risk_analysis_results.json"
    payload = {key: [] for key in ("course_overview", "risk_distribution", "at_risk_students", "intervention_plan", "student_risk_tracker")}
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    excel_path = root / "Student_Risk_Analysis.xlsx"
    book = load_workbook(excel_path)
    for sheet in book.worksheets:
        if sheet.max_row > 1:
            sheet.delete_rows(2, sheet.max_row - 1)
    book.save(excel_path)
    book.close()
    word_path = root / "Intervention_Plan.docx"
    document = Document(word_path)
    for table in document.tables:
        for row in table.rows[1:]:
            for cell in row.cells:
                cell.text = ""
    for paragraph in document.paragraphs:
        if paragraph.style.name == "Normal":
            paragraph.text = "Source snapshot received; detailed retention analysis is incomplete."
    document.save(word_path)
    (root / "Student_Risk_Tracker.md").write_text("# Student Risk Tracker\n\nThe detailed course tracker is incomplete.\n", encoding="utf-8")
    (root / "advisor_email.txt").write_text(
        "To: academic_advisors@university.edu\nSubject: Student Retention Risk Analysis - Action Required\n\nThe detailed analysis is incomplete.\n",
        encoding="utf-8",
    )


def _input_only(root: Path) -> None:
    for path in list(root.iterdir()):
        if path.is_file():
            path.unlink()
    for source in INPUTS.iterdir():
        if source.is_file():
            shutil.copy2(source, root / source.name)


PARTIALS: dict[str, tuple[Mutator, bool, bool]] = {
    "json_one_enrollment_wrong": (_json(lambda value: value["course_overview"][0].update({"enrollment_count": 2285})), False, False),
    "json_one_course_average_wrong": (_json(lambda value: value["course_overview"][1].update({"avg_score": 79.72})), False, False),
    "json_one_risk_count_wrong": (_json(lambda value: value["risk_distribution"][0].update({"student_count": 364})), False, False),
    "json_one_intervention_wrong": (_json(lambda value: value["intervention_plan"][1].update({"timeline": "Within 3 weeks"})), False, False),
    "json_one_tracker_level_wrong": (_json(lambda value: value["student_risk_tracker"][0].update({"Risk Level": "Medium"})), False, False),
    "scorer_threshold_shifted": (_scorer_threshold, True, False),
    "scorer_hardcoded_against_mutation": (_scorer_hardcoded, True, False),
    "excel_one_enrollment_wrong": (_excel("Course_Overview", "C2", 2285), False, False),
    "excel_one_average_wrong": (_excel("Course_Overview", "D3", 79.72), False, False),
    "excel_one_risk_count_wrong": (_excel("Risk_Distribution", "B2", 364), False, True),
    "excel_one_course_risk_wrong": (_excel("At_Risk_Students", "B2", 210), False, False),
    "excel_one_intervention_wrong": (_excel("Intervention_Plan", "B3", "Weekly office hours"), False, False),
    "excel_one_pct_formula_flattened": (_excel("Risk_Distribution", "C2", 9.77), False, True),
    "excel_one_pct_formula_wrong_ref": (_excel("Risk_Distribution", "C3", "=ROUND(B2/SUM($B$2:$B$4)*100,2)"), False, True),
    "excel_one_header_fill_missing": (_format_header, False, False),
    "word_one_heading_missing": (_word_heading, False, False),
    "word_one_course_average_wrong": (lambda root: _word_table(root, 0, "Основы финансов (Осень 2014)", 4, "79.72"), False, False),
    "word_one_risk_count_wrong": (lambda root: _word_table(root, 1, "High", 1, "364"), False, False),
    "word_one_risk_pct_wrong": (lambda root: _word_table(root, 1, "Medium", 2, "32.26%"), False, False),
    "word_one_intervention_wrong": (lambda root: _word_replace(root, "Group tutoring sessions", "Weekly office hours"), False, False),
    "word_source_boundary_missing": (lambda root: _word_replace(root, "Входы не содержат причин низкой успеваемости, результатов вмешательств или временного ряда; такие выводы не делались.", "Входы предоставлены университетом."), False, False),
    "word_unsupported_effect": (_word_unsupported, False, False),
    "tracker_h1_wrong": (_tracker_replace("# Student Risk Tracker", "# Retention Tracker"), False, False),
    "tracker_one_header_wrong": (_tracker_replace("Average Score", "Mean Score"), False, False),
    "tracker_one_course_count_wrong": (_tracker_replace("| Основы финансов (Осень 2013) | Low | 1872 |", "| Основы финансов (Осень 2013) | Low | 1873 |"), False, False),
    "email_wrong_recipient": (_email_replace("academic_advisors@university.edu", "student_success@university.edu"), False, False),
    "email_wrong_subject": (_email_replace("Student Retention Risk Analysis - Action Required", "Student Retention Summary"), False, False),
    "email_high_count_wrong": (_email_replace("363 студента High risk", "364 студента High risk"), False, False),
    "email_medium_action_wrong": (_email_replace("Group tutoring sessions", "Weekly office hours"), False, False),
    "moderate_risk_domain_conflict": (_moderate_risk_domain_conflict, False, True),
    "moderate_risk_domain_incomplete": (_moderate_risk_domain_incomplete, False, True),
    "severe_valid_but_incomplete_package": (_severe_valid_incomplete, False, True),
    "unexpected_debug_file": (lambda root: (root / "debug-notes.txt").write_text("unrequested\n", encoding="utf-8"), False, False),
}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    input_dir = JUDGE._input_dir()
    truth_with_details = JUDGE.V1.derive_truth(input_dir, include_student_details=True)
    truth = JUDGE.V1.public_truth(truth_with_details)
    golden_canonical = JUDGE._canonical(GOLDEN / "risk_scorer.py", input_dir, truth)
    golden_mutation = JUDGE._mutation(GOLDEN / "risk_scorer.py", input_dir, truth_with_details)
    golden_render = JUDGE._render(GOLDEN)
    expected_pct = [row["pct"] for row in truth["risk_distribution"]]
    original_canonical, original_mutation, original_render = JUDGE._canonical, JUDGE._mutation, JUDGE._render
    original_recalculate = JUDGE.V1.recalculate_excel
    assert JUDGE.evaluate(GOLDEN)["reward"] == 1.0
    records: list[dict[str, object]] = []

    def evaluate_case(case: Path, real_scorer: bool, real_recalc: bool) -> float:
        JUDGE._canonical = original_canonical if real_scorer else (lambda *_args: golden_canonical)
        JUDGE._mutation = original_mutation if real_scorer else (lambda *_args: golden_mutation)
        JUDGE._render = lambda _artifact: golden_render
        JUDGE.V1.recalculate_excel = original_recalculate if real_recalc else (lambda _path, _target: (expected_pct, "cached Golden recalculation"))
        try:
            return float(JUDGE.evaluate(case)["reward"])
        finally:
            JUDGE._canonical, JUDGE._mutation, JUDGE._render = original_canonical, original_mutation, original_render
            JUDGE.V1.recalculate_excel = original_recalculate

    with tempfile.TemporaryDirectory(prefix="retention-v2-negatives-") as raw:
        base = Path(raw)
        for name in REQUIRED:
            for mode in ("missing", "empty", "corrupt"):
                case = base / f"fatal-{mode}-{name}"
                _copy(case)
                target = case / name
                if mode == "missing":
                    target.unlink()
                elif mode == "empty":
                    target.write_bytes(b"")
                else:
                    target.write_bytes(b"not-a-valid-task-artifact")
                reward = evaluate_case(case, False, False)
                assert reward == 0, (mode, name, reward)
                records.append({"case": f"{mode}:{name}", "layer": "fatal_zero", "reward": reward})
        empty = base / "fatal-empty"
        empty.mkdir()
        reward = evaluate_case(empty, False, False)
        assert reward == 0
        records.append({"case": "empty-output", "layer": "fatal_zero", "reward": reward})
        input_only = base / "fatal-input-only"
        _copy(input_only)
        _input_only(input_only)
        reward = evaluate_case(input_only, False, False)
        assert reward == 0
        records.append({"case": "unmodified-input-only", "layer": "fatal_zero", "reward": reward})
        for name, (mutator, real_scorer, real_recalc) in PARTIALS.items():
            case = base / f"partial-{name}"
            _copy(case)
            mutator(case)
            reward = evaluate_case(case, real_scorer, real_recalc)
            assert 0 < reward < 1, (name, reward)
            records.append({"case": name, "layer": "partial", "reward": reward})
    partial = [float(record["reward"]) for record in records if record["layer"] == "partial"]
    report = {
        "total": len(records),
        "fatal_zero": sum(record["layer"] == "fatal_zero" for record in records),
        "partial_continuous": len(partial),
        "partial_reward_range": [min(partial), max(partial)],
        "cases": records,
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
