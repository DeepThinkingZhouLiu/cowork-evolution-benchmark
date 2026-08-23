#!/usr/bin/env python3
"""隐藏评测侧参考解法：根据任务规格生成标准办公交付物。"""

from __future__ import annotations

import argparse
import html
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt as PptPt

SPEC = json.loads(r'''{"required_inputs": ["项目预算_v1.xlsx", "项目计划_v1.docx", "项目状态_v1.pptx", "变更单.json"], "workbooks": [{"filename": "项目预算_v2.xlsx", "sheets": [{"name": "预算", "rows": [["项目", "金额_元", "负责人", "上线日期", "变更单"], ["客户数据平台", 920000, "林岚", "2026-10-20", "CR-2026-018"]]}, {"name": "变更记录", "rows": [["变更单", "变更原因", "原预算", "新预算"], ["CR-2026-018", "新增历史数据迁移与安全复核", 850000, 920000]]}]}], "documents": [{"filename": "项目计划_v2.docx", "title": "客户数据平台项目计划 v2", "sections": [{"heading": "变更摘要", "paragraphs": ["CR-2026-018：总预算调整为 920,000 元，负责人调整为林岚，上线日期调整为 2026-10-20。原因是新增历史数据迁移与安全复核。"]}, {"heading": "范围（保持不变）", "bullets": ["客户主数据", "标签服务", "权限审计"]}, {"heading": "版本控制", "paragraphs": ["本文件替代 v1，所有后续沟通使用 v2 基线。"]}]}], "presentations": [{"filename": "项目状态_v2.pptx", "slides": [{"title": "客户数据平台项目状态 v2", "metrics": [["总预算", "920,000 元"], ["负责人", "林岚"], ["上线日期", "2026-10-20"]]}, {"title": "本次变更", "bullets": ["变更单 CR-2026-018", "新增历史数据迁移与安全复核", "原范围保持不变"]}, {"title": "保持不变的范围", "bullets": ["客户主数据", "标签服务", "权限审计"]}]}]}''')


def font(run, size=10, bold=None):
    run.font.name = "Noto Sans CJK SC"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def build_docx(output, spec):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
    sec.left_margin, sec.right_margin = Cm(2.2), Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = "Noto Sans CJK SC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    normal.font.size = Pt(10.5)
    for name, size, color in (("Title", 22, "17365D"), ("Heading 1", 15, "17365D"), ("Heading 2", 12, "2F5597")):
        style = doc.styles[name]
        style.font.name = "Noto Sans CJK SC"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run(spec["title"]), 22, True)
    if spec.get("subtitle"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(p.add_run(spec["subtitle"]), 10)
    for section in spec["sections"]:
        doc.add_heading(section["heading"], level=section.get("level", 1))
        for text in section.get("paragraphs", []):
            p = doc.add_paragraph()
            font(p.add_run(str(text)), 10)
        for text in section.get("bullets", []):
            p = doc.add_paragraph(style="List Bullet")
            font(p.add_run(str(text)), 10)
        if section.get("table"):
            table_spec = section["table"]
            table = doc.add_table(rows=1, cols=len(table_spec["headers"]))
            table.style = "Table Grid"
            for col, value in enumerate(table_spec["headers"]):
                table.rows[0].cells[col].text = str(value)
                for run in table.rows[0].cells[col].paragraphs[0].runs:
                    font(run, 9, True)
            for values in table_spec["rows"]:
                cells = table.add_row().cells
                for col, value in enumerate(values):
                    cells[col].text = str(value)
                    for run in cells[col].paragraphs[0].runs:
                        font(run, 9)
    for change in spec.get("redlines", []):
        p = doc.add_paragraph()
        font(p.add_run(change.get("label", "修订") + "："), 10, True)
        deletion = OxmlElement("w:del")
        deletion.set(qn("w:author"), "CoworkEvoBench Reviewer")
        deletion.set(qn("w:date"), "2026-08-24T00:00:00Z")
        run = OxmlElement("w:r")
        text = OxmlElement("w:delText")
        text.text = change["old"]
        run.append(text)
        deletion.append(run)
        insertion = OxmlElement("w:ins")
        insertion.set(qn("w:author"), "CoworkEvoBench Reviewer")
        insertion.set(qn("w:date"), "2026-08-24T00:00:00Z")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = change["new"]
        run.append(text)
        insertion.append(run)
        p._p.append(deletion)
        p._p.append(insertion)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run("CoworkEvoBench 公开合成样例"), 8)
    doc.save(output)


def build_xlsx(output, spec):
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_spec in spec["sheets"]:
        ws = wb.create_sheet(sheet_spec["name"])
        for row in sheet_spec.get("rows", []):
            ws.append(row)
        for coordinate, value in sheet_spec.get("cells", {}).items():
            ws[coordinate] = value
        for coordinate, value in sheet_spec.get("formulas", {}).items():
            ws[coordinate] = value
        if ws.max_row:
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="17365D")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center")
        for column in ws.columns:
            width = max(len(str(cell.value or "")) for cell in column) + 2
            ws.column_dimensions[column[0].column_letter].width = min(max(width, 11), 34)
        ws.freeze_panes = "A2"
        if ws.max_row > 1 and ws.max_column > 1:
            ws.auto_filter.ref = ws.dimensions
        chart_spec = sheet_spec.get("chart")
        if chart_spec:
            chart = LineChart() if chart_spec.get("type") == "line" else BarChart()
            chart.title = chart_spec.get("title", "指标概览")
            data = Reference(ws, min_col=chart_spec.get("data_col", 2), min_row=1, max_row=ws.max_row)
            categories = Reference(ws, min_col=chart_spec.get("category_col", 1), min_row=2, max_row=ws.max_row)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(categories)
            ws.add_chart(chart, chart_spec.get("anchor", "H2"))
    wb.save(output)


def textbox(slide, text, x, y, w, h, size):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.name = "Noto Sans CJK SC"
    p.font.size = PptPt(size)
    p.font.color.rgb = PptRGBColor(33, 43, 54)
    return box


def build_pptx(output, spec):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for index, slide_spec in enumerate(spec["slides"], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = PptRGBColor(246, 248, 251)
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.16))
        band.fill.solid()
        band.fill.fore_color.rgb = PptRGBColor(23, 54, 93)
        band.line.fill.background()
        title = textbox(slide, slide_spec["title"], 0.7, 0.45, 11.8, 0.65, 25)
        title.text_frame.paragraphs[0].font.bold = True
        title.text_frame.paragraphs[0].font.color.rgb = PptRGBColor(23, 54, 93)
        y = 1.35
        for bullet in slide_spec.get("bullets", []):
            textbox(slide, "• " + str(bullet), 0.9, y, 11.4, 0.48, 15)
            y += 0.62
        metrics = slide_spec.get("metrics", [])
        for metric_index, metric in enumerate(metrics):
            x = 0.8 + metric_index * 3.8
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2), Inches(3.4), Inches(1.8))
            card.fill.solid()
            card.fill.fore_color.rgb = PptRGBColor(217, 234, 247)
            card.line.color.rgb = PptRGBColor(91, 155, 213)
            frame = card.text_frame
            frame.clear()
            p = frame.paragraphs[0]
            p.text = str(metric[1])
            p.alignment = PP_ALIGN.CENTER
            p.font.size = PptPt(25)
            p.font.bold = True
            p.font.color.rgb = PptRGBColor(23, 54, 93)
            p = frame.add_paragraph()
            p.text = str(metric[0])
            p.alignment = PP_ALIGN.CENTER
            p.font.size = PptPt(11)
        if slide_spec.get("table"):
            table_spec = slide_spec["table"]
            rows = [table_spec["headers"], *table_spec["rows"]]
            shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(0.8), Inches(1.55), Inches(11.7), Inches(min(4.8, 0.48 * len(rows))))
            for row_index, values in enumerate(rows):
                for col_index, value in enumerate(values):
                    cell = shape.table.cell(row_index, col_index)
                    cell.text = str(value)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = PptRGBColor(23, 54, 93) if row_index == 0 else PptRGBColor(255, 255, 255)
                    p = cell.text_frame.paragraphs[0]
                    p.font.size = PptPt(10)
                    p.font.color.rgb = PptRGBColor(255, 255, 255) if row_index == 0 else PptRGBColor(33, 43, 54)
                    p.font.bold = row_index == 0
        footer = textbox(slide, f"CoworkEvoBench 公开合成样例  |  {index}", 8.5, 7.0, 4.2, 0.25, 8)
        footer.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    prs.save(output)


def build_html(output, spec):
    title = html.escape(str(spec["title"]))
    subtitle = html.escape(str(spec.get("subtitle", "")))
    decision = html.escape(str(spec["decision"]))
    guardrail = html.escape(str(spec["guardrail"]))
    review_date = html.escape(str(spec["review_date"]))
    metrics = "".join(
        '<article class="metric"><strong>{}</strong><span>{}</span></article>'.format(
            html.escape(str(value)), html.escape(str(label))
        )
        for label, value in spec.get("metrics", [])
    )
    evidence = "".join(
        f"<li>{html.escape(str(item))}</li>" for item in spec.get("evidence", [])
    )
    actions = "".join(
        "<tr><td>{}</td><td>{}</td><td>{}</td></tr>".format(
            *(html.escape(str(value)) for value in row)
        )
        for row in spec.get("actions", [])
    )
    document = f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{title}</title>
  <style>
    :root {{ color-scheme: light; font-family: "Noto Sans CJK SC", sans-serif; }}
    body {{ margin: 0; background: #f4f7fb; color: #1f2937; }}
    main {{ max-width: 1040px; margin: 0 auto; padding: 40px 24px 64px; }}
    header {{ color: white; background: #17365d; border-radius: 18px; padding: 28px; }}
    h1 {{ margin: 0 0 8px; font-size: 30px; }}
    .metrics {{ display: grid; grid-template-columns: repeat(3, 1fr); gap: 16px; margin: 22px 0; }}
    .metric, section {{ background: white; border-radius: 14px; padding: 20px; box-shadow: 0 4px 16px #17365d12; }}
    .metric strong {{ display: block; color: #17365d; font-size: 28px; }}
    .metric span {{ color: #52606d; }}
    section {{ margin-top: 18px; }}
    .decision {{ border-left: 6px solid #16a085; }}
    .guardrail {{ border-left: 6px solid #e67e22; }}
    table {{ width: 100%; border-collapse: collapse; }}
    th, td {{ text-align: left; padding: 10px; border-bottom: 1px solid #dbe3ec; }}
    button {{ border: 0; border-radius: 9px; padding: 10px 14px; background: #2f5597; color: white; cursor: pointer; }}
    [hidden] {{ display: none; }}
    @media (max-width: 720px) {{ .metrics {{ grid-template-columns: 1fr; }} }}
  </style>
</head>
<body>
<main>
  <header><h1>{title}</h1><p>{subtitle}</p></header>
  <div class="metrics">{metrics}</div>
  <section class="decision"><h2>决策建议</h2><p>{decision}</p></section>
  <section class="guardrail"><h2>强制护栏</h2><p>{guardrail}</p><p>下一次复审：{review_date}</p></section>
  <section>
    <h2>研究证据</h2>
    <button id="evidence-button" type="button" aria-expanded="false" onclick="toggleEvidence()">展开研究证据</button>
    <ul id="evidence-list" hidden>{evidence}</ul>
  </section>
  <section><h2>扩展行动</h2><table><thead><tr><th>行动</th><th>责任</th><th>时间</th></tr></thead><tbody>{actions}</tbody></table></section>
</main>
<script>
function toggleEvidence() {{
  const list = document.getElementById('evidence-list');
  const button = document.getElementById('evidence-button');
  list.hidden = !list.hidden;
  button.setAttribute('aria-expanded', String(!list.hidden));
  button.textContent = list.hidden ? '展开研究证据' : '收起研究证据';
}}
</script>
</body>
</html>
"""
    output.write_text(document, encoding="utf-8")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    missing = [name for name in SPEC["required_inputs"] if not (args.input_dir / name).exists()]
    if missing:
        raise SystemExit(f"缺少输入文件：{missing}")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for spec in SPEC.get("documents", []):
        build_docx(args.output_dir / spec["filename"], spec)
    for spec in SPEC.get("workbooks", []):
        build_xlsx(args.output_dir / spec["filename"], spec)
    for spec in SPEC.get("presentations", []):
        build_pptx(args.output_dir / spec["filename"], spec)
    for spec in SPEC.get("html_pages", []):
        build_html(args.output_dir / spec["filename"], spec)


if __name__ == "__main__":
    main()
