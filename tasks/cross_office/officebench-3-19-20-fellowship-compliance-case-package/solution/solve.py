#!/usr/bin/env python3
"""OfficeBench 3-19/3-20 奖学金合规案件包的人工真值参考解法。"""

from __future__ import annotations

import argparse
import hashlib
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


PDF_NAME = "fellowship_application.pdf"
POLICY_NAME = "school_policy.jpg"
XLSX_NAME = "Fellowship_Compliance_Audit.xlsx"
DOCX_NAME = "Fellowship_Compliance_Decision.docx"
EML_NAME = "Fellowship_Application_Forward.eml"
ICS_NAME = "Fellowship_Compliance_Followup.ics"

SOURCE_HASHES = {
    PDF_NAME: "f317d69fbd1a208308bda56ba34322fecadb8564cbb5b4121c050314b0c6d180",
    POLICY_NAME: "b1d5b4d4b39028a8d40cac036055bfbbc6740b00d0a7cfe1c1406f52fd4a5f42",
}
FIXED_ZIP_TIME = (2026, 9, 1, 7, 14, 0)

NAVY = "17324D"
TEAL = "007C83"
BLUE = "DDEBF7"
GREEN = "E2F0D9"
AMBER = "FFF2CC"
RED = "FCE4D6"
WHITE = "FFFFFF"
LINE = "B8C7D9"
GRAY = "667085"


POLICY_ROWS = [
    ("P01", "The applicant's name must be provided.", "Alice Johnson is present in the Name field.", "PASS", "Yes", "Directly verifiable from the PDF."),
    ("P02", "The applicant's email must be a valid email address.", "alice.johnson@example.com has a valid address format.", "PASS", "Yes", "Format validation only; mailbox ownership is not inferred."),
    ("P03", "The application must include a statement of purpose with a minimum 100 words.", "Statement of Purpose contains 105 English words.", "PASS", "Yes", "Count uses contiguous English-letter word tokens."),
    ("P04", "The applicant must provide two references.", "One reference is present: Dr. John Smith, Professor, University X.", "FAIL", "Yes", "One additional reference is required."),
    ("P05", "All application materials must be submitted before the deadline.", "Neither a policy deadline nor a submission timestamp is provided.", "NOT_VERIFIABLE", "Yes", "Do not infer on-time or late submission."),
    ("P06", "Incomplete applications will not be considered.", "The application is incomplete because only one of two required references is present.", "FAIL", "Yes", "This converts the missing reference into an overall NO decision."),
]


@dataclass(frozen=True)
class ApplicationEvidence:
    name: str
    email: str
    statement: str
    statement_words: int
    references: tuple[str, ...]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def validate_inputs(input_dir: Path) -> None:
    for filename, expected in SOURCE_HASHES.items():
        path = input_dir / filename
        if not path.is_file():
            raise FileNotFoundError(f"缺少冻结输入：{filename}")
        actual = sha256(path)
        if actual != expected:
            raise ValueError(f"输入哈希不符：{filename} {actual}")


def extract_pdf_text(path: Path) -> str:
    process = subprocess.run(
        ["pdftotext", "-layout", str(path), "-"],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        timeout=120,
    )
    return process.stdout


def parse_application(path: Path) -> ApplicationEvidence:
    text = extract_pdf_text(path)
    name_match = re.search(r"Name:\s+([^\n]+)", text)
    email_match = re.search(r"Email:\s+([^\n]+)", text)
    statement_match = re.search(r"Statement of Purpose:\s*(.*?)\s*References:", text, flags=re.DOTALL)
    reference_match = re.search(r"References:\s*(.*?)(?:\f|\Z)", text, flags=re.DOTALL)
    if not all((name_match, email_match, statement_match, reference_match)):
        raise ValueError("无法从冻结 PDF 恢复申请字段")
    name = " ".join(name_match.group(1).split())
    email = " ".join(email_match.group(1).split())
    statement = " ".join(statement_match.group(1).split())
    words = re.findall(r"[A-Za-z]+(?:[-'’][A-Za-z]+)*", statement)
    references = tuple(
        " ".join(line.split())
        for line in reference_match.group(1).splitlines()
        if line.strip()
    )
    evidence = ApplicationEvidence(name, email, statement, len(words), references)
    expected = ApplicationEvidence(
        "Alice Johnson",
        "alice.johnson@example.com",
        statement,
        105,
        ("Dr. John Smith, Professor, University X",),
    )
    if evidence.name != expected.name or evidence.email != expected.email:
        raise ValueError("申请人字段与人工真值不符")
    if evidence.statement_words != expected.statement_words or evidence.references != expected.references:
        raise ValueError("陈述或参考人与人工真值不符")
    return evidence


def normalize_ooxml(path: Path) -> None:
    """统一 OOXML ZIP 顺序、时间戳和核心属性，使空目录回放可逐字节比较。"""

    with tempfile.TemporaryDirectory(prefix="fellowship-ooxml-") as temporary:
        target = Path(temporary) / path.name
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as destination:
            for name in sorted(source.namelist()):
                data = source.read(name)
                if name == "docProps/core.xml":
                    try:
                        root = etree.fromstring(data)
                        namespaces = {
                            "dcterms": "http://purl.org/dc/terms/",
                            "dc": "http://purl.org/dc/elements/1.1/",
                        }
                        for tag in ("created", "modified"):
                            nodes = root.xpath(f"//dcterms:{tag}", namespaces=namespaces)
                            for node in nodes:
                                node.text = "2026-09-01T07:14:00Z"
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except Exception:
                        pass
                if name.startswith("xl/charts/chart") and name.endswith(".xml"):
                    try:
                        root = etree.fromstring(data)
                        chart_namespace = "http://schemas.openxmlformats.org/drawingml/2006/chart"
                        id_nodes = root.xpath("//c:axId", namespaces={"c": chart_namespace})
                        stable_ids: dict[str, str] = {}
                        for node in id_nodes:
                            old_id = node.get("val")
                            if old_id is not None and old_id not in stable_ids:
                                stable_ids[old_id] = str(10_000_001 + len(stable_ids))
                        for node in root.xpath("//c:axId | //c:crossAx", namespaces={"c": chart_namespace}):
                            old_id = node.get("val")
                            if old_id in stable_ids:
                                node.set("val", stable_ids[old_id])
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except Exception:
                        pass
                info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                destination.writestr(info, data)
        shutil.copy2(target, path)


def recalculate_workbook(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="fellowship-recalc-") as temporary:
        output_dir = Path(temporary) / "output"
        profile_dir = Path(temporary) / "profile"
        output_dir.mkdir()
        process = subprocess.run(
            [
                "libreoffice",
                "--headless",
                f"-env:UserInstallation=file://{profile_dir}",
                "--convert-to",
                "xlsx",
                "--outdir",
                str(output_dir),
                str(path),
            ],
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            timeout=300,
        )
        converted = output_dir / path.name
        if process.returncode != 0 or not converted.is_file():
            raise RuntimeError(f"LibreOffice 公式重算失败：{process.stdout[-1000:]}")
        shutil.copy2(converted, path)


def restore_frozen_rows(path: Path, sheet_count: int) -> None:
    """LibreOffice 二次保存后恢复每张工作表的首行冻结。"""

    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with tempfile.TemporaryDirectory(prefix="fellowship-freeze-") as temporary:
        target = Path(temporary) / path.name
        sheet_names = {f"xl/worksheets/sheet{index}.xml" for index in range(1, sheet_count + 1)}
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w") as destination:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename in sheet_names:
                    root = etree.fromstring(data)
                    sheet_view = root.find(f"{{{namespace}}}sheetViews/{{{namespace}}}sheetView")
                    if sheet_view is not None:
                        for existing in sheet_view.findall(f"{{{namespace}}}pane"):
                            sheet_view.remove(existing)
                        pane = etree.Element(f"{{{namespace}}}pane")
                        pane.set("ySplit", "1")
                        pane.set("topLeftCell", "A2")
                        pane.set("activePane", "bottomLeft")
                        pane.set("state", "frozen")
                        sheet_view.insert(0, pane)
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                destination.writestr(info, data)
        shutil.copy2(target, path)


def add_table(sheet, name: str, reference: str) -> None:
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def style_sheet(sheet, widths: dict[str, float], orientation: str = "landscape") -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 34
    thin = Side(style="thin", color=LINE)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=9, color=NAVY)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.page_setup.orientation = orientation
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.4
    sheet.page_margins.bottom = 0.4
    sheet.oddFooter.center.text = "Fellowship compliance audit | &P / &N"
    sheet.oddFooter.center.size = 8
    sheet.oddFooter.center.color = GRAY


def build_workbook(evidence: ApplicationEvidence, output_path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.properties.creator = "EvoCowork OfficeBench adapter"
    workbook.properties.title = "Fellowship Compliance Audit"

    policy = workbook.create_sheet("Policy_Criteria")
    policy.append(["Criterion_ID", "Policy_Text", "Application_Evidence", "Status", "Determinative", "Notes"])
    for row_index, row in enumerate(POLICY_ROWS, start=2):
        values = list(row)
        if row_index == 2:
            values[3] = '=IF(Application_Evidence!B2<>"","PASS","FAIL")'
        elif row_index == 3:
            values[3] = '=IF(Application_Evidence!B4="TRUE","PASS","FAIL")'
        elif row_index == 4:
            values[3] = '=IF(Application_Evidence!B5>=100,"PASS","FAIL")'
        elif row_index == 5:
            values[3] = '=IF(Application_Evidence!B6>=2,"PASS","FAIL")'
        elif row_index == 7:
            values[3] = '=IF(COUNTIF(D2:D6,"FAIL")>0,"FAIL",IF(COUNTIF(D2:D6,"NOT_VERIFIABLE")>0,"NOT_VERIFIABLE","PASS"))'
        policy.append(values)
    style_sheet(policy, {"A": 12, "B": 49, "C": 55, "D": 19, "E": 16, "F": 42})
    for row in range(2, 8):
        policy.row_dimensions[row].height = 50
    policy.conditional_formatting.add("D2:D7", FormulaRule(formula=['D2="PASS"'], fill=PatternFill("solid", fgColor=GREEN)))
    policy.conditional_formatting.add("D2:D7", FormulaRule(formula=['D2="FAIL"'], fill=PatternFill("solid", fgColor=RED)))
    policy.conditional_formatting.add("D2:D7", FormulaRule(formula=['D2="NOT_VERIFIABLE"'], fill=PatternFill("solid", fgColor=AMBER)))
    add_table(policy, "FellowshipPolicyCriteriaTable", "A1:F7")
    policy.print_area = "A1:F7"

    application = workbook.create_sheet("Application_Evidence")
    application.append(["Evidence_Field", "Extracted_Value", "Source", "Validation_Note"])
    rows = [
        ("Applicant_Name", evidence.name, "fellowship_application.pdf / Name", "Provided"),
        ("Applicant_Email", evidence.email, "fellowship_application.pdf / Email", "Address format only"),
        ("Email_Format_Valid", "TRUE", "Regex structure check", "Mailbox ownership not inferred"),
        ("Statement_Word_Count", evidence.statement_words, "fellowship_application.pdf / Statement of Purpose", "Minimum is 100"),
        ("Reference_Count", len(evidence.references), "fellowship_application.pdf / References", "Policy requires 2"),
        ("Reference_1", evidence.references[0], "fellowship_application.pdf / References", "Only listed reference"),
        ("Policy_Deadline", "Not provided", "school_policy.jpg", "Timeliness cannot be verified"),
        ("Submission_Timestamp", "Not provided", "fellowship_application.pdf", "Timeliness cannot be verified"),
        ("Statement_of_Purpose", evidence.statement, "fellowship_application.pdf", "105-word source text"),
    ]
    for row in rows:
        application.append(row)
    style_sheet(application, {"A": 27, "B": 75, "C": 48, "D": 42})
    for row in range(2, 11):
        application.row_dimensions[row].height = 31 if row < 10 else 125
    add_table(application, "FellowshipApplicationEvidenceTable", "A1:D10")
    application.print_area = "A1:D10"

    summary = workbook.create_sheet("Decision_Summary")
    summary.append(["Metric", "Value", "Evidence_or_Rule"])
    summary_rows = [
        ("PASS_Count", '=COUNTIF(Policy_Criteria!D2:D7,"PASS")', "Policy_Criteria status count"),
        ("FAIL_Count", '=COUNTIF(Policy_Criteria!D2:D7,"FAIL")', "Policy_Criteria status count"),
        ("NOT_VERIFIABLE_Count", '=COUNTIF(Policy_Criteria!D2:D7,"NOT_VERIFIABLE")', "Policy_Criteria status count"),
        ("Overall_Decision", '=IF(B3>0,"NO",IF(B4>0,"NOT_VERIFIABLE","YES"))', "Any FAIL -> NO; otherwise any evidence gap -> NOT_VERIFIABLE"),
        ("Missing_Reference_Count", "=MAX(0,2-Application_Evidence!B6)", "Two required minus extracted count"),
        ("Statement_Word_Count", "=Application_Evidence!B5", "Minimum 100 words"),
        ("Officer_Address_Status", "BLOCKED_MISSING_OFFICER_ADDRESS", "No officer email appears in either source"),
        ("Meeting_Start_UTC", "2024-05-31 13:00", "OfficeBench 3-20 upstream evaluator"),
        ("Meeting_Duration", "Not provided", "ICS intentionally omits DTEND and DURATION"),
    ]
    for row in summary_rows:
        summary.append(row)
    summary["E1"] = "Status"
    summary["F1"] = "Count"
    summary["E2"], summary["F2"] = "PASS", "=B2"
    summary["E3"], summary["F3"] = "FAIL", "=B3"
    summary["E4"], summary["F4"] = "NOT_VERIFIABLE", "=B4"
    style_sheet(summary, {"A": 35, "B": 34, "C": 68, "D": 4, "E": 24, "F": 14, "G": 3, "H": 16, "I": 16, "J": 16})
    for row in range(2, 11):
        summary.row_dimensions[row].height = 26
    summary["B5"].fill = PatternFill("solid", fgColor=RED)
    summary["B5"].font = Font(name="Aptos Display", size=14, bold=True, color="9C0006")
    add_table(summary, "FellowshipDecisionSummaryTable", "A1:C10")
    chart = BarChart()
    chart.style = 10
    chart.title = "Policy status counts"
    chart.y_axis.title = "Criteria"
    chart.x_axis.title = "Status"
    chart.height = 7.0
    chart.width = 12.5
    chart.add_data(Reference(summary, min_col=6, min_row=1, max_row=4), titles_from_data=True)
    chart.set_categories(Reference(summary, min_col=5, min_row=2, max_row=4))
    summary.add_chart(chart, "H2")
    summary.print_area = "A1:N10"

    delivery = workbook.create_sheet("Delivery_Control")
    delivery.append(["Deliverable", "Purpose", "Decision_Consistency", "Delivery_Status", "Control_Note"])
    delivery_rows = [
        (XLSX_NAME, "Auditable policy and evidence ledger", "Overall decision must be NO", "READY", "Native formulas and chart retained"),
        (DOCX_NAME, "Human-readable compliance decision", "Must match Excel policy statuses", "READY", "No unsupported timeliness claim"),
        (EML_NAME, "Officer forwarding draft with original application", "Body must report NO and evidence gaps", "BLOCKED", "Replace reserved .invalid address before sending"),
        (ICS_NAME, "Tentative compliance follow-up", "Description must report NO and evidence gaps", "TENTATIVE", "No duration, location or attendee address supplied"),
    ]
    for row in delivery_rows:
        delivery.append(row)
    style_sheet(delivery, {"A": 46, "B": 42, "C": 43, "D": 20, "E": 52})
    for row in range(2, 6):
        delivery.row_dimensions[row].height = 42
    delivery.conditional_formatting.add("D2:D5", FormulaRule(formula=['D2="READY"'], fill=PatternFill("solid", fgColor=GREEN)))
    delivery.conditional_formatting.add("D2:D5", FormulaRule(formula=['OR(D2="BLOCKED",D2="TENTATIVE")'], fill=PatternFill("solid", fgColor=AMBER)))
    add_table(delivery, "FellowshipDeliveryControlTable", "A1:E5")
    delivery.print_area = "A1:E5"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    recalculate_workbook(output_path)
    restore_frozen_rows(output_path, 4)
    normalize_ooxml(output_path)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_width(cell, centimeters: float) -> None:
    properties = cell._tc.get_or_add_tcPr()
    width = properties.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        properties.append(width)
    width.set(qn("w:w"), str(int(centimeters * 567)))
    width.set(qn("w:type"), "dxa")


def add_docx_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float], font_size: float = 7.6) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_width(cell, widths[index])
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(font_size)
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(WHITE)
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.text = value
            set_cell_width(cell, widths[index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            if row_index % 2 == 0:
                set_cell_shading(cell, "F5F8FB")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = RGBColor.from_string(NAVY)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Heading 1")
    paragraph.add_run(text)


def add_body(document: Document, text: str, bold_prefix: str | None = None) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.05
    if bold_prefix and text.startswith(bold_prefix):
        paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)


def add_page_number(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.extend([field_begin, instruction, field_end])
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)


def build_document(evidence: ApplicationEvidence, output_path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9)
    styles["Normal"].font.color.rgb = RGBColor.from_string(NAVY)
    styles["Heading 1"].font.name = "Aptos Display"
    styles["Heading 1"].font.size = Pt(14)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].font.color.rgb = RGBColor.from_string(TEAL)
    styles["Heading 1"].paragraph_format.space_before = Pt(8)
    styles["Heading 1"].paragraph_format.space_after = Pt(4)
    document.core_properties.title = "Fellowship Application Compliance Decision"
    document.core_properties.author = "EvoCowork OfficeBench adapter"

    header = section.header.paragraphs[0]
    header.text = "FELLOWSHIP COMPLIANCE | EVIDENCE-BOUND REVIEW"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(TEAL)
    add_page_number(section)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(2)
    title_run = title.add_run("Fellowship Application\nCompliance Decision")
    title_run.font.name = "Aptos Display"
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = document.add_paragraph("Alice Johnson | OfficeBench 3-19 / 3-20 evidence package")
    subtitle.paragraph_format.space_after = Pt(12)
    for run in subtitle.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(GRAY)

    decision = document.add_table(rows=1, cols=2)
    decision.autofit = False
    decision.alignment = WD_TABLE_ALIGNMENT.LEFT
    labels = [("OVERALL DECISION", "NO"), ("PRIMARY REASON", "Only 1 of 2 required references is present")]
    for index, (label, value) in enumerate(labels):
        cell = decision.rows[0].cells[index]
        set_cell_width(cell, 8.2)
        set_cell_shading(cell, RED if index == 0 else BLUE)
        paragraph = cell.paragraphs[0]
        paragraph.add_run(label + "\n").bold = True
        value_run = paragraph.add_run(value)
        value_run.bold = True
        value_run.font.size = Pt(17 if index == 0 else 10)
        value_run.font.color.rgb = RGBColor.from_string("9C0006" if index == 0 else NAVY)

    add_heading(document, "Decision")
    add_body(document, "Decision: NO. The application is incomplete because the policy requires two references and the PDF contains one. The statement contains 105 words and meets the 100-word minimum.")
    add_body(document, "Timeliness is NOT_VERIFIABLE: the policy image gives no deadline date and the application PDF gives no submission timestamp. This evidence gap is not treated as proof of on-time or late submission.")

    add_heading(document, "Applicant Evidence")
    add_docx_table(
        document,
        ["Field", "Extracted value", "Assessment"],
        [
            ["Applicant", evidence.name, "Provided"],
            ["Email", evidence.email, "Valid address format; mailbox ownership not inferred"],
            ["Statement of Purpose", f"{evidence.statement_words} words", "PASS; minimum is 100"],
            ["References", f"{len(evidence.references)}: {evidence.references[0]}", "FAIL; policy requires 2"],
            ["Policy deadline", "Not provided", "NOT_VERIFIABLE"],
            ["Submission timestamp", "Not provided", "NOT_VERIFIABLE"],
        ],
        [3.4, 8.1, 5.4],
    )

    add_heading(document, "Policy-by-Policy Assessment")
    add_docx_table(
        document,
        ["ID", "Policy", "Status", "Evidence"],
        [[row[0], row[1], row[3], row[2]] for row in POLICY_ROWS],
        [1.1, 6.7, 2.6, 6.5],
        font_size=7.1,
    )

    add_heading(document, "Evidence Gaps")
    add_body(document, "The source package does not provide a policy deadline, a submission timestamp, or an officer email address. It also does not provide meeting duration, location, or an attendee mailbox. These values remain explicitly unavailable across every deliverable.")

    document.add_page_break()
    add_heading(document, "Required Remediation")
    add_docx_table(
        document,
        ["Priority", "Action", "Reason", "Owner / dependency"],
        [
            ["1", "Obtain one additional reference", "Two references are required; only one is present", "Applicant"],
            ["2", "Obtain the policy deadline and submission evidence", "Timeliness cannot currently be verified", "Program office"],
            ["3", "Replace the reserved .invalid recipient before sending", "No officer email address is supplied", "Program office"],
            ["4", "Confirm duration/location only if the office supplies them", "The upstream meeting request specifies only a start time", "Meeting organizer"],
        ],
        [1.3, 6.4, 6.2, 3.1],
        font_size=7.2,
    )

    add_heading(document, "Delivery Controls")
    add_docx_table(
        document,
        ["Artifact", "Control", "Status"],
        [
            [XLSX_NAME, "Formula-backed policy counts and NO decision", "READY"],
            [DOCX_NAME, "Matches Excel and states evidence boundaries", "READY"],
            [EML_NAME, "Original PDF attached; recipient address unresolved", "BLOCKED"],
            [ICS_NAME, "13:00 UTC start only; no invented duration/location/attendee", "TENTATIVE"],
        ],
        [7.0, 7.3, 2.7],
    )
    add_body(document, "Control conclusion: the compliance decision is NO. The email remains blocked until a real officer address is supplied, while the follow-up calendar item remains tentative.")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    normalize_ooxml(output_path)


def build_email(evidence: ApplicationEvidence, pdf_path: Path, output_path: Path) -> None:
    message = EmailMessage(policy=SMTP)
    message["From"] = "Bob <sender-address-unavailable@example.invalid>"
    message["To"] = "Fellowship Officer <address-unavailable@example.invalid>"
    message["Subject"] = "Fellowship application compliance review - Alice Johnson"
    message["Message-ID"] = "<officebench-3-19-20-fellowship@evocowork.invalid>"
    message["X-EvoCowork-Delivery-Status"] = "BLOCKED_MISSING_OFFICER_ADDRESS"
    message["X-EvoCowork-Decision"] = "NO"
    body = (
        "Fellowship Officer,\n\n"
        "Decision: NO. Alice Johnson's application contains one reference, while the school policy requires two. "
        f"The Statement of Purpose contains {evidence.statement_words} words and meets the 100-word minimum.\n\n"
        "Evidence gap: neither a specific policy deadline nor a submission timestamp is provided, so timeliness is NOT_VERIFIABLE. "
        "No conclusion of on-time or late submission has been made.\n\n"
        "Delivery control: this draft is blocked because the officer's real email address is not provided. "
        "Replace the reserved example.invalid address before sending. The original fellowship_application.pdf is attached unchanged.\n"
    )
    message.set_content(body, subtype="plain", charset="utf-8")
    message.add_attachment(
        pdf_path.read_bytes(),
        maintype="application",
        subtype="pdf",
        filename=PDF_NAME,
        disposition="attachment",
    )
    message.set_boundary("===============EvoCoworkOfficeBench31920==")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("wb") as handle:
        BytesGenerator(handle, policy=SMTP.clone(max_line_length=78), mangle_from_=False).flatten(message)


def fold_ics_line(line: str, limit: int = 73) -> list[str]:
    """按 UTF-8 字节安全折行；本题字段为 ASCII，但仍保留通用边界。"""

    chunks: list[str] = []
    remaining = line
    first = True
    while len(remaining.encode("utf-8")) > limit:
        cut = limit
        while len(remaining[:cut].encode("utf-8")) > limit:
            cut -= 1
        chunks.append(("" if first else " ") + remaining[:cut])
        remaining = remaining[cut:]
        first = False
        limit = 72
    chunks.append(("" if first else " ") + remaining)
    return chunks


def build_calendar(output_path: Path) -> None:
    description = (
        "Decision: NO for Alice Johnson; the Statement of Purpose has 105 words, but only one of two required references is present. "
        "Timeliness is NOT_VERIFIABLE because the policy deadline and submission timestamp are not provided. "
        "Officer email, meeting duration, location, and attendee address are unavailable."
    )
    description = description.replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")
    logical_lines = [
        "BEGIN:VCALENDAR",
        "VERSION:2.0",
        "PRODID:-//EvoCowork//OfficeBench 3-19 3-20 Fellowship Compliance//EN",
        "CALSCALE:GREGORIAN",
        "METHOD:PUBLISH",
        "BEGIN:VEVENT",
        "UID:officebench-3-19-20-fellowship-followup@evocowork.invalid",
        "DTSTAMP:20260901T071400Z",
        "DTSTART:20240531T130000Z",
        "SUMMARY:Fellowship compliance follow-up with officer",
        "STATUS:TENTATIVE",
        f"DESCRIPTION:{description}",
        "X-EVOCOWORK-DECISION:NO",
        "X-EVOCOWORK-EVIDENCE-GAP:DEADLINE_AND_SUBMISSION_TIMESTAMP_UNAVAILABLE",
        "END:VEVENT",
        "END:VCALENDAR",
    ]
    physical_lines = [chunk for line in logical_lines for chunk in fold_ics_line(line)]
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_bytes(("\r\n".join(physical_lines) + "\r\n").encode("utf-8"))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    validate_inputs(args.input_dir)
    evidence = parse_application(args.input_dir / PDF_NAME)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for existing in args.output_dir.iterdir():
        if existing.is_file() and existing.name in {XLSX_NAME, DOCX_NAME, EML_NAME, ICS_NAME}:
            existing.unlink()
    build_workbook(evidence, args.output_dir / XLSX_NAME)
    build_document(evidence, args.output_dir / DOCX_NAME)
    build_email(evidence, args.input_dir / PDF_NAME, args.output_dir / EML_NAME)
    build_calendar(args.output_dir / ICS_NAME)


if __name__ == "__main__":
    main()
