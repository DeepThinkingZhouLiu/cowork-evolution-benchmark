#!/usr/bin/env python3
"""奖学金合规案件包的 v2 原子连续评分；v1 保存在 judge_v1.py。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from email import policy
from email.parser import BytesParser
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from openpyxl import load_workbook


TASK_ID = "cowork-evo/officebench-3-19-20-fellowship-compliance-case-package"
RUBRIC_VERSION = "officebench-3-19-20-fellowship-compliance-v0.2"
JUDGE_VERSION = "fellowship-compliance-atomic-package-v2.0"
XLSX = "Fellowship_Compliance_Audit.xlsx"
DOCX = "Fellowship_Compliance_Decision.docx"
EML = "Fellowship_Application_Forward.eml"
ICS = "Fellowship_Compliance_Followup.ics"
REQUIRED = (XLSX, DOCX, EML, ICS)
POSITIVE_WEIGHTS = {"R001":6,"R002":12,"R003":10,"R004":9,"R005":8,"R006":12,"R007":10,"R008":9,"R009":7,"R010":7,"R011":4,"R012":6}
PENALTY_WEIGHTS = {"P001":-12,"P002":-4}


def _load_v1():
    path = Path(__file__).with_name("judge_v1.py")
    spec = importlib.util.spec_from_file_location("fellowship_compliance_judge_v1", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v1 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


V1 = _load_v1()


POLICY_ROWS = [
    ("P01", V1.POLICY_TEXTS[0], "Alice Johnson is present in the Name field.", "PASS", "Yes", "Directly verifiable from the PDF."),
    ("P02", V1.POLICY_TEXTS[1], "alice.johnson@example.com has a valid address format.", "PASS", "Yes", "Format validation only; mailbox ownership is not inferred."),
    ("P03", V1.POLICY_TEXTS[2], "Statement of Purpose contains 105 English words.", "PASS", "Yes", "Count uses contiguous English-letter word tokens."),
    ("P04", V1.POLICY_TEXTS[3], "One reference is present: Dr. John Smith, Professor, University X.", "FAIL", "Yes", "One additional reference is required."),
    ("P05", V1.POLICY_TEXTS[4], "Neither a policy deadline nor a submission timestamp is provided.", "NOT_VERIFIABLE", "Yes", "Do not infer on-time or late submission."),
    ("P06", V1.POLICY_TEXTS[5], "The application is incomplete because only one of two required references is present.", "FAIL", "Yes", "This converts the missing reference into an overall NO decision."),
]
SUMMARY_ROWS = [
    ("PASS_Count", 3, "Policy_Criteria status count"),
    ("FAIL_Count", 2, "Policy_Criteria status count"),
    ("NOT_VERIFIABLE_Count", 1, "Policy_Criteria status count"),
    ("Overall_Decision", "NO", "Any FAIL -> NO; otherwise any evidence gap -> NOT_VERIFIABLE"),
    ("Missing_Reference_Count", 1, "Two required minus extracted count"),
    ("Statement_Word_Count", 105, "Minimum 100 words"),
    ("Officer_Address_Status", "BLOCKED_MISSING_OFFICER_ADDRESS", "No officer email appears in either source"),
    ("Meeting_Start_UTC", "2024-05-31 13:00", "OfficeBench 3-20 upstream evaluator"),
    ("Meeting_Duration", "Not provided", "ICS intentionally omits DTEND and DURATION"),
]
DELIVERY_ROWS = [
    (XLSX, "Auditable policy and evidence ledger", "Overall decision must be NO", "READY", "Native formulas and chart retained"),
    (DOCX, "Human-readable compliance decision", "Must match Excel policy statuses", "READY", "No unsupported timeliness claim"),
    (EML, "Officer forwarding draft with original application", "Body must report NO and evidence gaps", "BLOCKED", "Replace reserved .invalid address before sending"),
    (ICS, "Tentative compliance follow-up", "Description must report NO and evidence gaps", "TENTATIVE", "No duration, location or attendee address supplied"),
]


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _same(actual: Any, expected: Any) -> bool:
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        try:
            return abs(float(actual) - float(expected)) <= 0.011
        except (TypeError, ValueError):
            return False
    return V1.normalize(actual) == V1.normalize(expected)


def _ratio(findings: Iterable[tuple[bool, str]]) -> tuple[float, list[str], int, int]:
    items = list(findings)
    earned = sum(bool(ok) for ok, _ in items)
    total = len(items)
    return earned / max(1, total), [label for ok, label in items if not ok], earned, total


def _criterion(identifier: str, score: float, failures: list[str], earned: float, total: float) -> dict[str, Any]:
    return {"criterion_id":identifier,"score":round(max(0.0,min(1.0,score)),6),"evidence":"全部任务专属原子检查通过" if not failures else "未通过："+"；".join(failures),"method":"deterministic","raw":{"earned_units":round(earned,6),"total_units":round(total,6)}}


def _task_root() -> Path:
    return Path(__file__).resolve().parents[1]


def _input_dir() -> Path:
    return Path("/app/input_files") if Path("/app/input_files").is_dir() else _task_root()/"data"/"input_files"


def _core(artifact: Path) -> tuple[bool,list[str]]:
    if not artifact.is_dir():
        return False,["输出目录不存在"]
    errors=[]
    for name in REQUIRED:
        path=artifact/name
        if not path.is_file() or path.stat().st_size==0:
            errors.append(f"核心交付缺失或为空：{name}");continue
        try:
            if name==XLSX:
                if not zipfile.is_zipfile(path):raise ValueError("不是有效 OOXML ZIP")
                book=load_workbook(path,read_only=True,data_only=False)
                if not book.sheetnames:raise ValueError("工作簿无工作表")
                book.close()
            elif name==DOCX:
                if not zipfile.is_zipfile(path):raise ValueError("不是有效 OOXML ZIP")
                document=Document(path)
                if not any(p.text.strip() for p in document.paragraphs):raise ValueError("Word 无正文")
            elif name==EML:
                message=BytesParser(policy=policy.default).parse(path.open("rb"))
                if not message.get("MIME-Version") or message.get_body(preferencelist=("plain",)) is None:raise ValueError("不是完整 MIME 邮件")
            else:
                text=path.read_text(encoding="utf-8")
                if "BEGIN:VCALENDAR" not in text or "END:VCALENDAR" not in text:raise ValueError("不是完整 iCalendar")
        except Exception as exc:
            errors.append(f"核心交付无法解析：{name}: {type(exc).__name__}: {exc}")
    source_errors=[]
    for name,expected in V1.SOURCE_HASHES.items():
        source=_input_dir()/name
        if not source.is_file() or _sha256(source)!=expected:source_errors.append(name)
    if source_errors:errors.append(f"冻结输入缺失或哈希错误：{source_errors}")
    return not errors,errors


def _fatal(errors:list[str],hashes:list[str])->dict[str,Any]:
    return {"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":hashes or ["0"*64],"criterion_results":[_criterion(key,0,errors,0,0) for key in [*POSITIVE_WEIGHTS,*PENALTY_WEIGHTS]],"reward":0.0,"passed":False,"hurdle_failures":["R001"],"penalties":[]}


def _recalculate(path:Path)->tuple[dict[str,list[list[Any]]]|None,str]:
    with tempfile.TemporaryDirectory(prefix="fellowship-v2-recalc-") as raw:
        base=Path(raw);source=base/"source";output=base/"output";profile=base/"profile";source.mkdir();output.mkdir();copied=source/path.name;shutil.copy2(path,copied)
        env=os.environ.copy();env["HOME"]=str(base/"home");env["NO_PROXY"]="localhost,127.0.0.1,::1"
        try:completed=subprocess.run(["libreoffice","--headless",f"-env:UserInstallation=file://{profile}","--convert-to","xlsx","--outdir",str(output),str(copied)],stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,timeout=300,check=False,env=env)
        except Exception as exc:return None,str(exc)
        result=output/path.name
        if completed.returncode!=0 or not result.is_file():return None,f"退出={completed.returncode}，文件={result.is_file()}，日志={completed.stdout[-500:]}"
        try:
            book=load_workbook(result,data_only=True)
            values={sheet.title:[[cell.value for cell in row] for row in sheet.iter_rows()] for sheet in book.worksheets};book.close();return values,completed.stdout[-500:]
        except Exception as exc:return None,str(exc)


def _value(values:dict[str,list[list[Any]]]|None,sheet:str,row:int,column:int)->Any:
    try:return values[sheet][row-1][column-1] if values is not None else None
    except (KeyError,IndexError):return None


def _rows(book:Any,sheet_name:str,columns:int)->list[list[Any]]:
    if sheet_name not in book.sheetnames:return []
    sheet=book[sheet_name];result=[]
    for row in range(2,sheet.max_row+1):
        values=[sheet.cell(row,column).value for column in range(1,columns+1)]
        if any(value not in (None,"") for value in values):result.append(values)
    return result


def _excel_atomic(artifact:Path,truth:dict[str,Any])->dict[str,Any]:
    path=artifact/XLSX;book=load_workbook(path,data_only=False);values,recalc_log=_recalculate(path)
    policy_findings=[];app_findings=[];formula_findings=[];delivery_findings=[];native_findings=[]
    expected_sheets=["Policy_Criteria","Application_Evidence","Decision_Summary","Delivery_Control"]
    for index,name in enumerate(expected_sheets):
        actual=book.sheetnames[index] if index<len(book.sheetnames) else None
        delivery_findings.append((actual==name,f"工作表第{index+1}项错误：{actual!r}"))
    policy_sheet=book["Policy_Criteria"] if "Policy_Criteria" in book.sheetnames else None
    policy_headers=[policy_sheet.cell(1,c).value for c in range(1,7)] if policy_sheet else []
    for index,expected in enumerate(V1.POLICY_HEADERS):policy_findings.append((_same(policy_headers[index] if index<len(policy_headers) else None,expected),f"Policy_Criteria 表头第{index+1}列错误"))
    policy_rows=_rows(book,"Policy_Criteria",6)
    policy_findings.append((len(policy_rows)==6,f"政策行数错误：{len(policy_rows)}"))
    for row_index,expected in enumerate(POLICY_ROWS,2):
        actual=policy_rows[row_index-2] if row_index-2<len(policy_rows) else []
        for column,value in enumerate(expected,1):
            found=_value(values,"Policy_Criteria",row_index,column) if column==4 else (actual[column-1] if column-1<len(actual) else None)
            policy_findings.append((_same(found,value),f"Policy_Criteria {expected[0]} 第{column}列错误"))
    statement=str(truth.get("statement", ""))
    application_rows=[
        ("Applicant_Name",truth.get("name"),"fellowship_application.pdf / Name","Provided"),
        ("Applicant_Email",truth.get("email"),"fellowship_application.pdf / Email","Address format only"),
        ("Email_Format_Valid","TRUE","Regex structure check","Mailbox ownership not inferred"),
        ("Statement_Word_Count",105,"fellowship_application.pdf / Statement of Purpose","Minimum is 100"),
        ("Reference_Count",1,"fellowship_application.pdf / References","Policy requires 2"),
        ("Reference_1","Dr. John Smith, Professor, University X","fellowship_application.pdf / References","Only listed reference"),
        ("Policy_Deadline","Not provided","school_policy.jpg","Timeliness cannot be verified"),
        ("Submission_Timestamp","Not provided","fellowship_application.pdf","Timeliness cannot be verified"),
        ("Statement_of_Purpose",statement,"fellowship_application.pdf","105-word source text"),
    ]
    app_sheet=book["Application_Evidence"] if "Application_Evidence" in book.sheetnames else None
    app_headers=[app_sheet.cell(1,c).value for c in range(1,5)] if app_sheet else []
    for index,expected in enumerate(V1.APPLICATION_HEADERS):app_findings.append((_same(app_headers[index] if index<len(app_headers) else None,expected),f"Application_Evidence 表头第{index+1}列错误"))
    actual_app=_rows(book,"Application_Evidence",4);app_findings.append((len(actual_app)==9,f"申请证据行数错误：{len(actual_app)}"))
    for row_index,expected in enumerate(application_rows):
        actual=actual_app[row_index] if row_index<len(actual_app) else []
        for column,value in enumerate(expected):app_findings.append((_same(actual[column] if column<len(actual) else None,value),f"Application_Evidence {expected[0]} 第{column+1}列错误"))
    expected_formula_tokens={
        "D2":("APPLICATION_EVIDENCE!B2","PASS","FAIL"),"D3":("APPLICATION_EVIDENCE!B4","PASS","FAIL"),"D4":("APPLICATION_EVIDENCE!B5",">=100","PASS","FAIL"),"D5":("APPLICATION_EVIDENCE!B6",">=2","PASS","FAIL"),"D7":("COUNTIF","D2:D6","FAIL","NOT_VERIFIABLE"),
        "B2":("COUNTIF","POLICY_CRITERIA!D2:D7","PASS"),"B3":("COUNTIF","POLICY_CRITERIA!D2:D7","FAIL"),"B4":("COUNTIF","POLICY_CRITERIA!D2:D7","NOT_VERIFIABLE"),"B5":("B3","B4","NO","NOT_VERIFIABLE","YES"),"B6":("MAX","APPLICATION_EVIDENCE!B6"),"B7":("APPLICATION_EVIDENCE!B5",),
    }
    for cell,tokens in expected_formula_tokens.items():
        sheet=policy_sheet if cell.startswith("D") else (book["Decision_Summary"] if "Decision_Summary" in book.sheetnames else None)
        formula=sheet[cell].value if sheet else None;normalized=str(formula or "").upper().replace("$","").replace(" ","")
        formula_findings.append((isinstance(formula,str) and formula.startswith("="),f"{sheet.title if sheet else '?'}!{cell} 不是公式"))
        for token in tokens:formula_findings.append((token.upper().replace("$","").replace(" ","") in normalized,f"{sheet.title if sheet else '?'}!{cell} 缺公式要素 {token}"))
    for row,expected in enumerate(V1.POLICY_STATUSES,2):formula_findings.append((_same(_value(values,"Policy_Criteria",row,4),expected),f"Policy_Criteria!D{row} 重算错误"))
    for row,expected in enumerate((3,2,1,"NO",1,105),2):formula_findings.append((_same(_value(values,"Decision_Summary",row,2),expected),f"Decision_Summary!B{row} 重算错误：{recalc_log}"))
    summary_sheet=book["Decision_Summary"] if "Decision_Summary" in book.sheetnames else None
    summary_headers=[summary_sheet.cell(1,c).value for c in range(1,3+1)] if summary_sheet else []
    for index,expected in enumerate(V1.SUMMARY_HEADERS):delivery_findings.append((_same(summary_headers[index] if index<len(summary_headers) else None,expected),f"Decision_Summary 表头第{index+1}列错误"))
    actual_summary=_rows(book,"Decision_Summary",3);delivery_findings.append((len(actual_summary)==9,f"Decision_Summary 行数错误：{len(actual_summary)}"))
    for row_index,expected in enumerate(SUMMARY_ROWS,2):
        actual=actual_summary[row_index-2] if row_index-2<len(actual_summary) else []
        for column,value in enumerate(expected,1):
            found=_value(values,"Decision_Summary",row_index,column) if column==2 and row_index<=7 else (actual[column-1] if column-1<len(actual) else None)
            delivery_findings.append((_same(found,value),f"Decision_Summary {expected[0]} 第{column}列错误"))
    delivery_sheet=book["Delivery_Control"] if "Delivery_Control" in book.sheetnames else None
    delivery_headers=[delivery_sheet.cell(1,c).value for c in range(1,6)] if delivery_sheet else []
    for index,expected in enumerate(V1.DELIVERY_HEADERS):delivery_findings.append((_same(delivery_headers[index] if index<len(delivery_headers) else None,expected),f"Delivery_Control 表头第{index+1}列错误"))
    actual_delivery=_rows(book,"Delivery_Control",5);delivery_findings.append((len(actual_delivery)==4,f"Delivery_Control 行数错误：{len(actual_delivery)}"))
    for row_index,expected in enumerate(DELIVERY_ROWS):
        actual=actual_delivery[row_index] if row_index<len(actual_delivery) else []
        for column,value in enumerate(expected):delivery_findings.append((_same(actual[column] if column<len(actual) else None,value),f"Delivery_Control {expected[0]} 第{column+1}列错误"))
    for sheet_name in expected_sheets:
        if sheet_name not in book.sheetnames:
            native_findings.extend([(False,f"缺工作表 {sheet_name}")]*5);continue
        sheet=book[sheet_name]
        native_findings.extend([(len(sheet.tables)>=1,f"{sheet_name} 无原生表"),(bool(sheet.auto_filter.ref),f"{sheet_name} 无筛选"),(str(sheet.freeze_panes or "")=="A2",f"{sheet_name} 未冻结首行"),(all(bool(c.font.bold) for c in sheet[1] if c.value is not None),f"{sheet_name} 表头未加粗"),(len(sheet._images)==0,f"{sheet_name} 用图片替代内容")])
    native_findings.extend([(sum(len(s._charts) for s in book.worksheets)==1,"政策状态图数量不是1"),(sum(len(s.conditional_formatting) for s in book.worksheets)>=2,"条件格式区域少于2")])
    text=" ".join(V1.normalize(_value(values,s.title,r,c)) for s in book.worksheets for r in range(1,s.max_row+1) for c in range(1,s.max_column+1) if _value(values,s.title,r,c) is not None)
    book.close()
    return {"policy":_ratio(policy_findings),"application":_ratio(app_findings),"formula":_ratio(formula_findings),"delivery":_ratio(delivery_findings),"native":native_findings,"text":text,"policy_rows":policy_rows,"application_rows":actual_app,"summary_rows":actual_summary,"delivery_rows":actual_delivery}


def _table_rows(document:Document,index:int)->list[list[str]]:
    if index>=len(document.tables):return []
    return [[V1.normalize(cell.text) for cell in row.cells] for row in document.tables[index].rows]


def _word_atomic(artifact:Path)->tuple[float,list[str],int,int,str,list[list[list[str]]],list[str],int]:
    document=Document(artifact/DOCX);paragraphs=[V1.normalize(p.text) for p in document.paragraphs if V1.normalize(p.text)];headings=[V1.normalize(p.text) for p in document.paragraphs if p.style and p.style.name.startswith("Heading")];tables=[_table_rows(document,i) for i in range(len(document.tables))];text=" ".join(paragraphs+[cell for table in tables for row in table for cell in row])
    findings=[]
    findings.append(("Fellowship Application Compliance Decision" in text,"Word 标题错误"))
    for heading in ("Decision","Applicant Evidence","Policy-by-Policy Assessment","Evidence Gaps","Required Remediation","Delivery Controls"):findings.append((heading in headings,f"Word 缺章节 {heading}"))
    for token in ("Decision: NO","only 1 of 2 required references","105 words","100-word minimum","NOT_VERIFIABLE","no deadline date","no submission timestamp","not treated as proof of on-time or late submission"):findings.append((token.casefold() in text.casefold(),f"Word 缺事实/边界 {token}"))
    evidence_expected=[("Applicant","Alice Johnson","Provided"),("Email","alice.johnson@example.com","Valid address format; mailbox ownership not inferred"),("Statement of Purpose","105 words","PASS; minimum is 100"),("References","1: Dr. John Smith, Professor, University X","FAIL; policy requires 2"),("Policy deadline","Not provided","NOT_VERIFIABLE"),("Submission timestamp","Not provided","NOT_VERIFIABLE")]
    evidence=tables[1][1:] if len(tables)>1 else [];findings.append((len(evidence)==6,f"Word 申请证据行数错误：{len(evidence)}"))
    for index,expected in enumerate(evidence_expected):
        actual=evidence[index] if index<len(evidence) else []
        for column,value in enumerate(expected):findings.append((_same(actual[column] if column<len(actual) else None,value),f"Word 申请证据第{index+1}行第{column+1}列错误"))
    policies=tables[2][1:] if len(tables)>2 else [];findings.append((len(policies)==6,f"Word 政策行数错误：{len(policies)}"))
    for index,expected in enumerate(POLICY_ROWS):
        wanted=(expected[0],expected[1],expected[3],expected[2]);actual=policies[index] if index<len(policies) else []
        for column,value in enumerate(wanted):findings.append((_same(actual[column] if column<len(actual) else None,value),f"Word 政策 {expected[0]} 第{column+1}列错误"))
    remediation=tables[3][1:] if len(tables)>3 else [];findings.append((len(remediation)==4,f"Word 整改行数错误：{len(remediation)}"))
    for token in ("Obtain one additional reference","Obtain the policy deadline and submission evidence","Replace the reserved .invalid recipient before sending","Confirm duration/location only if the office supplies them"):findings.append((any(token in cell for row in remediation for cell in row),f"Word 缺整改 {token}"))
    delivery=tables[4][1:] if len(tables)>4 else [];findings.append((len(delivery)==4,f"Word 交付控制行数错误：{len(delivery)}"))
    for index,(filename,_,_,status,_) in enumerate(DELIVERY_ROWS):
        actual=delivery[index] if index<len(delivery) else []
        findings.extend([(_same(actual[0] if actual else None,filename),f"Word 交付文件第{index+1}行错误"),(_same(actual[2] if len(actual)>2 else None,status),f"Word 交付状态 {filename} 错误")])
    with zipfile.ZipFile(artifact/DOCX) as archive:media=sum(1 for name in archive.namelist() if name.startswith("word/media/") and not name.endswith("/"))
    score,failures,earned,total=_ratio(findings);return score,failures,earned,total,text,tables,headings,media


def _email_atomic(artifact:Path)->tuple[float,list[str],int,int,str,str]:
    message=BytesParser(policy=policy.default).parse((artifact/EML).open("rb"));body_part=message.get_body(preferencelist=("plain",));body=V1.normalize(body_part.get_content()) if body_part else "";attachments=list(message.iter_attachments());recipient=V1.normalize(message.get("To"));findings=[(message.is_multipart(),"EML 不是 multipart"),(_same(message.get("Subject"),"Fellowship application compliance review - Alice Johnson"),"EML 主题错误"),(_same(recipient,"Fellowship Officer <address-unavailable@example.invalid>"),"EML 收件人错误"),(_same(message.get("X-EvoCowork-Delivery-Status"),"BLOCKED_MISSING_OFFICER_ADDRESS"),"EML 阻塞头错误"),(_same(message.get("X-EvoCowork-Decision"),"NO"),"EML 决策头错误")]
    for token in ("Decision: NO","Alice Johnson","one reference","requires two","105 words","NOT_VERIFIABLE","blocked","example.invalid","no conclusion of on-time or late submission"):findings.append((token.casefold() in body.casefold(),f"EML 正文缺 {token}"))
    findings.append((len(attachments)==1,f"EML 附件数错误：{len(attachments)}"))
    if attachments:
        attachment=attachments[0];payload=attachment.get_payload(decode=True) or b"";source=_input_dir()/V1.PDF_NAME
        findings.extend([(_same(attachment.get_filename(),V1.PDF_NAME),"EML 附件名错误"),(attachment.get_content_type()=="application/pdf","EML 附件 MIME 错误"),(payload==source.read_bytes(),"EML 附件与输入 PDF 字节不一致")])
    else:findings.extend([(False,"EML 缺附件名"),(False,"EML 缺附件 MIME"),(False,"EML 缺 PDF 字节")])
    score,failures,earned,total=_ratio(findings);return score,failures,earned,total,body,recipient


def _calendar_atomic(artifact:Path)->tuple[float,list[str],int,int,str,dict[str,list[str]]]:
    raw=(artifact/ICS).read_bytes();text=raw.decode("utf-8");lines=V1.unfold_ics(text);properties={};in_event=False
    for line in lines:
        if line=="BEGIN:VEVENT":in_event=True;continue
        if line=="END:VEVENT":in_event=False;continue
        if in_event and ":" in line:
            key,value=line.split(":",1);properties.setdefault(key.split(";",1)[0].upper(),[]).append(value)
    description=" ".join(properties.get("DESCRIPTION",[])).replace("\\;",";").replace("\\,",",")
    findings=[(b"\r\n" in raw and raw.replace(b"\r\n",b"").find(b"\n")<0,"ICS 不是严格 CRLF"),(lines.count("BEGIN:VEVENT")==1,"ICS VEVENT 数量错误"),(properties.get("UID")==["officebench-3-19-20-fellowship-followup@evocowork.invalid"],"ICS UID 错误"),(properties.get("DTSTART")==["20240531T130000Z"],"ICS DTSTART 错误"),(properties.get("SUMMARY")==["Fellowship compliance follow-up with officer"],"ICS SUMMARY 错误"),(properties.get("STATUS")==["TENTATIVE"],"ICS STATUS 错误"),(properties.get("X-EVOCOWORK-DECISION")==["NO"],"ICS 决策错误"),(properties.get("X-EVOCOWORK-EVIDENCE-GAP")==["DEADLINE_AND_SUBMISSION_TIMESTAMP_UNAVAILABLE"],"ICS 证据缺口头错误")]
    for token in ("Decision: NO","Alice Johnson","105 words","only one of two required references","NOT_VERIFIABLE","policy deadline","submission timestamp"):findings.append((token.casefold() in description.casefold(),f"ICS 描述缺 {token}"))
    for forbidden in ("DTEND","DURATION","LOCATION","ATTENDEE","ORGANIZER"):findings.append((forbidden not in properties,f"ICS 含未授权字段 {forbidden}"))
    score,failures,earned,total=_ratio(findings);return score,failures,earned,total," ".join(lines),properties


def _cross(excel_text:str,word_text:str,email_body:str,ics_text:str)->tuple[float,list[str],int,int]:
    findings=[]
    for token in ("Alice Johnson","NO","105","one","two","NOT_VERIFIABLE"):
        for label,text in (("Excel",excel_text),("Word",word_text),("EML",email_body),("ICS",ics_text)):findings.append((token.casefold() in text.casefold(),f"{label} 缺跨文件事实 {token}"))
    for label,text in (("Excel",excel_text),("Word",word_text),("EML",email_body)):findings.append(("BLOCKED" in text.upper() or ".invalid" in text,f"{label} 缺邮件阻塞"))
    return _ratio(findings)


def _boundary(excel_text:str,word_text:str,email_body:str,ics_text:str,recipient:str)->tuple[float,list[str],int,int]:
    combined=" ".join((excel_text,word_text,email_body,ics_text));findings=[("FAIL" in combined and "NOT_VERIFIABLE" in combined,"未区分 FAIL 与 NOT_VERIFIABLE"),("one" in combined.casefold() and "two" in combined.casefold(),"缺参考人不足边界"),("deadline" in combined.casefold() and "submission timestamp" in combined.casefold(),"缺时间证据边界"),(recipient=="Fellowship Officer <address-unavailable@example.invalid>","未使用保留收件人")]
    for label,text in (("Word",word_text),("EML",email_body),("ICS",ics_text)):findings.append(("NOT_VERIFIABLE" in text,f"{label} 未声明时间不可核验"))
    unsupported=V1.unsupported_claims([excel_text,word_text,email_body,ics_text],recipient);findings.append((not unsupported,f"编造={unsupported}"));return _ratio(findings)


def _render(artifact:Path)->tuple[float,list[str],int,int]:
    result=V1.render_office(artifact);return _ratio([(result.get("xlsx_pages")==4,f"Excel 渲染页数={result.get('xlsx_pages')}，错误={result.get('errors')}"),(result.get("docx_pages")==2,f"Word 渲染页数={result.get('docx_pages')}，错误={result.get('errors')}")])


def _unsupported(excel_text:str,word_text:str,email_body:str,ics_text:str,recipient:str)->tuple[float,list[str],float]:
    labels=V1.unsupported_claims([excel_text,word_text,email_body,ics_text],recipient);severity=min(1.0,0.45+0.15*len(labels)) if labels else 0.0;return severity,labels,1.0


def _scope(artifact:Path,packages:dict[str,dict[str,Any]])->tuple[float,list[str],float]:
    extras=sorted(p.name for p in artifact.iterdir() if p.name not in REQUIRED);danger=[]
    for name,package in packages.items():
        if package.get("broken") or package.get("macros") or package.get("external") or package.get("missing_targets"):danger.append(f"{name}:{package}")
    labels=([f"额外文件={extras}"] if extras else [])+danger;return (1.0 if labels else 0.0),labels,1.0


def _safe_positive(call,*args):
    try:return call(*args)
    except Exception as exc:return 0.0,[f"任务专属检查异常：{type(exc).__name__}: {exc}"],0,1


def _safe_penalty(call,*args):
    try:return call(*args)
    except Exception as exc:return 1.0,[f"严重边界无法核验：{type(exc).__name__}: {exc}"],1.0


def evaluate(artifact:Path)->dict[str,Any]:
    ok,errors=_core(artifact);hashes=[_sha256(artifact/name) for name in REQUIRED if (artifact/name).is_file()]
    if not ok:return _fatal(errors,hashes)
    truth,source_errors,_=V1.validate_sources(_task_root())
    if source_errors:return _fatal(source_errors,hashes)
    excel=_excel_atomic(artifact,truth);word=_safe_positive(_word_atomic,artifact);email=_safe_positive(_email_atomic,artifact);calendar=_safe_positive(_calendar_atomic,artifact)
    word_score,word_failures,word_earned,word_total,word_text,word_tables,word_headings,word_media=word
    email_score,email_failures,email_earned,email_total,email_body,recipient=email
    calendar_score,calendar_failures,calendar_earned,calendar_total,ics_text,_properties=calendar
    cross=_safe_positive(_cross,excel["text"],word_text,email_body,ics_text);boundary=_safe_positive(_boundary,excel["text"],word_text,email_body,ics_text,recipient)
    packages={name:V1.inspect_ooxml(artifact/name) for name in (XLSX,DOCX)}
    native_findings=[*excel["native"],(len(word_tables)>=5,"Word 原生表格少于5张"),(len(word_headings)>=6,"Word 原生标题少于6个"),(word_media==0,"Word 用图片替代内容")]
    native_findings.extend((not package.get("broken") and not package.get("macros") and not package.get("external") and not package.get("missing_targets"),f"{name} OOXML 不安全：{package}") for name,package in packages.items())
    native=_ratio(native_findings)
    render=_safe_positive(_render,artifact)
    checks={"R001":(1.0,[],len(REQUIRED),len(REQUIRED)),"R002":excel["policy"],"R003":excel["application"],"R004":excel["formula"],"R005":excel["delivery"],"R006":word[:4],"R007":email[:4],"R008":calendar[:4],"R009":cross,"R010":boundary,"R011":native,"R012":render}
    results=[];scores={}
    for key,(score,failures,earned,total) in checks.items():results.append(_criterion(key,score,failures,earned,total));scores[key]=score
    penalties={"P001":_safe_penalty(_unsupported,excel["text"],word_text,email_body,ics_text,recipient),"P002":_safe_penalty(_scope,artifact,packages)}
    for key,(severity,labels,total) in penalties.items():results.append(_criterion(key,severity,labels,severity*total,total));scores[key]=severity
    denominator=sum(POSITIVE_WEIGHTS.values());positive=sum(POSITIVE_WEIGHTS[k]*scores[k] for k in POSITIVE_WEIGHTS);deduction=sum(abs(PENALTY_WEIGHTS[k])*scores[k] for k in PENALTY_WEIGHTS);reward=round(max(0.0,min(1.0,(positive-deduction)/denominator)),6)
    return {"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":hashes,"criterion_results":results,"reward":reward,"passed":reward>=0.8,"hurdle_failures":[],"penalties":[{"id":k,"amount":round(abs(PENALTY_WEIGHTS[k])*scores[k]/denominator,6)} for k in PENALTY_WEIGHTS if scores[k]>0]}


def main()->None:
    parser=argparse.ArgumentParser();parser.add_argument("--artifact-dir",type=Path,required=True);parser.add_argument("--result",type=Path,required=True);parser.add_argument("--reward-file",type=Path,required=True);parser.add_argument("--scoring-version",choices=("v1","v2"),default="v2");args=parser.parse_args()
    try:result=V1.evaluate(args.artifact_dir) if args.scoring_version=="v1" else evaluate(args.artifact_dir)
    except Exception as exc:result=_fatal([f"Judge 异常：{type(exc).__name__}: {exc}"],[])
    rendered=json.dumps(result,ensure_ascii=False,indent=2);args.result.parent.mkdir(parents=True,exist_ok=True);args.reward_file.parent.mkdir(parents=True,exist_ok=True);args.result.write_text(rendered+"\n",encoding="utf-8");args.reward_file.write_text(f"{float(result['reward']):.6f}\n",encoding="utf-8");print(rendered)


if __name__=="__main__":main()
