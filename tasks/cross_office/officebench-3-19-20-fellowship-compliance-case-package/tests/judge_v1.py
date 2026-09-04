#!/usr/bin/env python3
"""OfficeBench 3-19/3-20 奖学金合规案件包的独立确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import posixpath
import re
import shutil
import subprocess
import tempfile
import zipfile
from email import policy
from email.parser import BytesParser
from pathlib import Path

from docx import Document
from lxml import etree
from openpyxl import load_workbook


TASK_ID = "cowork-evo/officebench-3-19-20-fellowship-compliance-case-package"
RUBRIC_VERSION = "officebench-3-19-20-fellowship-compliance-v0.1"
JUDGE_VERSION = "officebench-3-19-20-fellowship-compliance-deterministic-v0.1"
PDF_NAME = "fellowship_application.pdf"
POLICY_NAME = "school_policy.jpg"
XLSX_NAME = "Fellowship_Compliance_Audit.xlsx"
DOCX_NAME = "Fellowship_Compliance_Decision.docx"
EML_NAME = "Fellowship_Application_Forward.eml"
ICS_NAME = "Fellowship_Compliance_Followup.ics"
REQUIRED_FILES = [XLSX_NAME, DOCX_NAME, EML_NAME, ICS_NAME]
SHEETS = ["Policy_Criteria", "Application_Evidence", "Decision_Summary", "Delivery_Control"]
SOURCE_HASHES = {
    PDF_NAME: "f317d69fbd1a208308bda56ba34322fecadb8564cbb5b4121c050314b0c6d180",
    POLICY_NAME: "b1d5b4d4b39028a8d40cac036055bfbbc6740b00d0a7cfe1c1406f52fd4a5f42",
}
POLICY_TEXTS = [
    "The applicant's name must be provided.",
    "The applicant's email must be a valid email address.",
    "The application must include a statement of purpose with a minimum 100 words.",
    "The applicant must provide two references.",
    "All application materials must be submitted before the deadline.",
    "Incomplete applications will not be considered.",
]
POLICY_STATUSES = ["PASS", "PASS", "PASS", "FAIL", "NOT_VERIFIABLE", "FAIL"]
POLICY_HEADERS = ["Criterion_ID", "Policy_Text", "Application_Evidence", "Status", "Determinative", "Notes"]
APPLICATION_HEADERS = ["Evidence_Field", "Extracted_Value", "Source", "Validation_Note"]
SUMMARY_HEADERS = ["Metric", "Value", "Evidence_or_Rule"]
DELIVERY_HEADERS = ["Deliverable", "Purpose", "Decision_Consistency", "Delivery_Status", "Control_Note"]
SUSPICIOUS_MARKERS = ["solution/golden", "judge.py", "expectations.json", "judge-result", "/root/", "/home/", "c:\\"]

SPECS = {
    "R001": ("完整交付与冻结输入", 5, "hurdle", 0.0),
    "R002": ("六条政策和状态", 10, "hurdle", 0.3),
    "R003": ("申请证据真值", 10, "hurdle", 0.3),
    "R004": ("原生公式与重算", 8, "hurdle", 0.35),
    "R005": ("NO 决策与缺失参考人", 10, "hurdle", 0.4),
    "R006": ("时间证据边界", 8, "hurdle", 0.45),
    "R007": ("Excel 原生审计结构", 5, "positive", None),
    "R008": ("Word 结构与表格", 5, "hurdle", 0.5),
    "R009": ("Word 合规真值", 8, "hurdle", 0.45),
    "R010": ("EML 草稿与原始附件", 8, "hurdle", 0.4),
    "R011": ("ICS 复核事件", 6, "hurdle", 0.4),
    "R012": ("四产物一致性", 6, "hurdle", 0.4),
    "R013": ("可渲染与专业排版", 3, "positive", None),
    "R014": ("无政策、证据、公式或决策错误", -30, "penalty", None),
    "R015": ("无编造事实", -25, "penalty", None),
    "R016": ("无缺失、损坏、静态化或调试泄漏", -25, "penalty", None),
    "R017": ("EML 无附件或收件人错误", -20, "penalty", None),
    "R018": ("ICS 无改期或未授权字段", -20, "penalty", None),
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def normalize(value: object) -> str:
    return re.sub(r"\s+", " ", "" if value is None else str(value)).strip()


def source_path(task_root: Path, filename: str) -> Path | None:
    candidates = [
        Path("/app/input_files") / filename,
        task_root / "data" / "input_files" / filename,
        task_root / "environment" / "input_files" / filename,
    ]
    return next((candidate for candidate in candidates if candidate.is_file()), None)


def extract_pdf_truth(path: Path) -> tuple[dict[str, object], list[str]]:
    errors: list[str] = []
    try:
        process = subprocess.run(
            ["pdftotext", "-layout", str(path), "-"],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            timeout=120,
        )
        if process.returncode != 0:
            return {}, [f"pdftotext_return:{process.returncode}"]
        text = process.stdout
        name = normalize(re.search(r"Name:\s+([^\n]+)", text).group(1))  # type: ignore[union-attr]
        email = normalize(re.search(r"Email:\s+([^\n]+)", text).group(1))  # type: ignore[union-attr]
        statement = normalize(re.search(r"Statement of Purpose:\s*(.*?)\s*References:", text, flags=re.DOTALL).group(1))  # type: ignore[union-attr]
        reference_block = re.search(r"References:\s*(.*?)(?:\f|\Z)", text, flags=re.DOTALL).group(1)  # type: ignore[union-attr]
        references = [normalize(line) for line in reference_block.splitlines() if normalize(line)]
        words = re.findall(r"[A-Za-z]+(?:[-'’][A-Za-z]+)*", statement)
        truth = {
            "name": name,
            "email": email,
            "statement": statement,
            "statement_words": len(words),
            "references": references,
        }
        if truth != {
            "name": "Alice Johnson",
            "email": "alice.johnson@example.com",
            "statement": statement,
            "statement_words": 105,
            "references": ["Dr. John Smith, Professor, University X"],
        }:
            errors.append(f"unexpected_pdf_truth:{truth}")
        return truth, errors
    except Exception as exc:
        return {}, [f"pdf_parse:{type(exc).__name__}:{exc}"]


def validate_sources(task_root: Path) -> tuple[dict[str, object], list[str], dict[str, Path]]:
    errors: list[str] = []
    paths: dict[str, Path] = {}
    for filename, expected in SOURCE_HASHES.items():
        path = source_path(task_root, filename)
        if path is None:
            errors.append(f"missing_source:{filename}")
            continue
        paths[filename] = path
        actual = sha256(path)
        if actual != expected:
            errors.append(f"source_hash:{filename}:{actual}")
    truth: dict[str, object] = {}
    if PDF_NAME in paths:
        truth, pdf_errors = extract_pdf_truth(paths[PDF_NAME])
        errors.extend(pdf_errors)
    return truth, errors, paths


def relationship_source(rel_name: str) -> str:
    if rel_name == "_rels/.rels":
        return ""
    marker = "/_rels/"
    if marker not in rel_name or not rel_name.endswith(".rels"):
        return ""
    prefix, filename = rel_name.split(marker, 1)
    return posixpath.join(prefix, filename[:-5])


def inspect_ooxml(path: Path) -> dict[str, object]:
    result: dict[str, object] = {"broken": False, "macros": [], "external": [], "missing_targets": []}
    try:
        with zipfile.ZipFile(path, "r") as archive:
            names = set(archive.namelist())
            bad = archive.testzip()
            if bad:
                result["broken"] = True
                result["missing_targets"].append(f"bad_crc:{bad}")  # type: ignore[union-attr]
            result["macros"] = sorted(name for name in names if "vbaproject" in name.lower())
            for rel_name in sorted(name for name in names if name.endswith(".rels")):
                try:
                    root = etree.fromstring(archive.read(rel_name))
                except Exception as exc:
                    result["broken"] = True
                    result["missing_targets"].append(f"rels_parse:{rel_name}:{type(exc).__name__}")  # type: ignore[union-attr]
                    continue
                source = relationship_source(rel_name)
                base = posixpath.dirname(source)
                for relationship in root:
                    target = relationship.get("Target", "")
                    if relationship.get("TargetMode") == "External":
                        result["external"].append(f"{rel_name}:{target}")  # type: ignore[union-attr]
                        continue
                    if not target or target.startswith("#"):
                        continue
                    resolved = posixpath.normpath(posixpath.join(base, target.lstrip("/")))
                    if resolved.startswith("../") or resolved not in names:
                        result["missing_targets"].append(f"{rel_name}:{target}->{resolved}")  # type: ignore[union-attr]
    except Exception as exc:
        result["broken"] = True
        result["missing_targets"].append(f"zip:{type(exc).__name__}:{exc}")  # type: ignore[union-attr]
    return result


def headers(sheet, count: int) -> list[str]:
    return [normalize(sheet.cell(1, column).value) for column in range(1, count + 1)]


def formula_cell(cell) -> bool:
    return cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("="))


def workbook_audit(path: Path, truth: dict[str, object]) -> dict[str, object]:
    audit: dict[str, object] = {
        "errors": [], "policy": False, "application": False, "formulas": False,
        "decision": False, "boundary": False, "delivery": False, "native": {}, "text": "",
    }
    try:
        formulas = load_workbook(path, data_only=False)
        values = load_workbook(path, data_only=True)
    except Exception as exc:
        audit["errors"].append(f"open:{type(exc).__name__}:{exc}")  # type: ignore[union-attr]
        return audit
    audit["text"] = " ".join(
        normalize(cell.value)
        for sheet in values.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    if formulas.sheetnames != SHEETS or values.sheetnames != SHEETS:
        audit["errors"].append(f"sheets:{formulas.sheetnames}")  # type: ignore[union-attr]
        return audit
    policy_f, policy_v = formulas["Policy_Criteria"], values["Policy_Criteria"]
    app_f, app_v = formulas["Application_Evidence"], values["Application_Evidence"]
    summary_f, summary_v = formulas["Decision_Summary"], values["Decision_Summary"]
    delivery_f, delivery_v = formulas["Delivery_Control"], values["Delivery_Control"]

    policy_rows = [
        [normalize(policy_v.cell(row, column).value) for column in range(1, 7)]
        for row in range(2, 8)
    ]
    audit["policy"] = (
        headers(policy_f, 6) == POLICY_HEADERS
        and [row[0] for row in policy_rows] == [f"P{index:02d}" for index in range(1, 7)]
        and [row[1] for row in policy_rows] == POLICY_TEXTS
        and [row[3] for row in policy_rows] == POLICY_STATUSES
        and all(row[4] == "Yes" for row in policy_rows)
        and "105" in policy_rows[2][2]
        and "Dr. John Smith" in policy_rows[3][2]
        and "Neither a policy deadline nor a submission timestamp" in policy_rows[4][2]
    )

    app_map = {normalize(app_v.cell(row, 1).value): app_v.cell(row, 2).value for row in range(2, 11)}
    audit["application"] = (
        headers(app_f, 4) == APPLICATION_HEADERS
        and normalize(app_map.get("Applicant_Name")) == truth.get("name")
        and normalize(app_map.get("Applicant_Email")) == truth.get("email")
        and normalize(app_map.get("Email_Format_Valid")).upper() in {"TRUE", "1"}
        and app_map.get("Statement_Word_Count") == truth.get("statement_words")
        and app_map.get("Reference_Count") == len(truth.get("references", []))
        and normalize(app_map.get("Reference_1")) == truth.get("references", [""])[0]
        and normalize(app_map.get("Policy_Deadline")) == "Not provided"
        and normalize(app_map.get("Submission_Timestamp")) == "Not provided"
        and normalize(app_map.get("Statement_of_Purpose")) == truth.get("statement")
    )

    required_policy_formula_rows = [2, 3, 4, 5, 7]
    formula_checks = all(formula_cell(policy_f.cell(row, 4)) for row in required_policy_formula_rows)
    formula_checks = formula_checks and all(formula_cell(summary_f.cell(row, 2)) for row in range(2, 8))
    formula_checks = formula_checks and [summary_v.cell(row, 2).value for row in range(2, 8)] == [3, 2, 1, "NO", 1, 105]
    audit["formulas"] = formula_checks
    audit["decision"] = normalize(summary_v["B5"].value) == "NO" and summary_v["B3"].value == 2 and summary_v["B6"].value == 1
    audit["boundary"] = (
        summary_v["B4"].value == 1
        and normalize(summary_v["B8"].value) == "BLOCKED_MISSING_OFFICER_ADDRESS"
        and normalize(summary_v["B9"].value) == "2024-05-31 13:00"
        and normalize(summary_v["B10"].value) == "Not provided"
    )
    expected_delivery = [XLSX_NAME, DOCX_NAME, EML_NAME, ICS_NAME]
    audit["delivery"] = (
        headers(delivery_f, 5) == DELIVERY_HEADERS
        and [normalize(delivery_v.cell(row, 1).value) for row in range(2, 6)] == expected_delivery
        and [normalize(delivery_v.cell(row, 4).value) for row in range(2, 6)] == ["READY", "READY", "BLOCKED", "TENTATIVE"]
    )
    audit["native"] = {
        "tables": sum(len(sheet.tables) for sheet in formulas.worksheets),
        "charts": sum(len(sheet._charts) for sheet in formulas.worksheets),
        "freeze": all(sheet.freeze_panes == "A2" for sheet in formulas.worksheets),
        "conditional_format_ranges": sum(len(sheet.conditional_formatting) for sheet in formulas.worksheets),
        "conditional_format_rules": sum(
            len(rules)
            for sheet in formulas.worksheets
            for rules in sheet.conditional_formatting._cf_rules.values()
        ),
    }
    return audit


def document_audit(path: Path) -> dict[str, object]:
    result: dict[str, object] = {"errors": [], "text": "", "headings": [], "tables": 0, "structure": False, "truth": False, "boundary": False}
    try:
        document = Document(path)
    except Exception as exc:
        result["errors"].append(f"open:{type(exc).__name__}:{exc}")  # type: ignore[union-attr]
        return result
    paragraphs = [normalize(paragraph.text) for paragraph in document.paragraphs if normalize(paragraph.text)]
    table_text = [normalize(cell.text) for table in document.tables for row in table.rows for cell in row.cells if normalize(cell.text)]
    text = " ".join(paragraphs + table_text)
    headings = [normalize(paragraph.text) for paragraph in document.paragraphs if paragraph.style and paragraph.style.name.startswith("Heading")]
    required_headings = ["Decision", "Applicant Evidence", "Policy-by-Policy Assessment", "Evidence Gaps", "Required Remediation", "Delivery Controls"]
    result.update({"text": text, "headings": headings, "tables": len(document.tables)})
    result["structure"] = (
        "Fellowship Application Compliance Decision" in text
        and all(heading in headings for heading in required_headings)
        and len(document.tables) >= 5
    )
    truth_tokens = [
        "Decision: NO", "Alice Johnson", "alice.johnson@example.com", "105 words",
        "Dr. John Smith, Professor, University X", "only 1 of 2 required references",
        "PASS", "FAIL", "NOT_VERIFIABLE",
    ]
    result["truth"] = all(token.lower() in text.lower() for token in truth_tokens)
    boundary_tokens = [
        "no deadline date", "no submission timestamp", "not treated as proof of on-time or late submission",
        "no officer email address", "does not provide meeting duration, location, or an attendee mailbox",
        "reserved .invalid recipient",
    ]
    result["boundary"] = all(token.lower() in text.lower() for token in boundary_tokens)
    return result


def email_audit(path: Path, source_pdf: Path | None) -> dict[str, object]:
    result: dict[str, object] = {"errors": [], "text": "", "headers": False, "body": False, "attachment": False, "safe_recipient": False}
    try:
        message = BytesParser(policy=policy.default).parse(path.open("rb"))
        body_part = message.get_body(preferencelist=("plain",))
        body = normalize(body_part.get_content()) if body_part else ""
        attachments = list(message.iter_attachments())
        result["text"] = body
        result["headers"] = (
            normalize(message.get("Subject")) == "Fellowship application compliance review - Alice Johnson"
            and normalize(message.get("X-EvoCowork-Delivery-Status")) == "BLOCKED_MISSING_OFFICER_ADDRESS"
            and normalize(message.get("X-EvoCowork-Decision")) == "NO"
            and message.is_multipart()
        )
        result["safe_recipient"] = normalize(message.get("To")) == "Fellowship Officer <address-unavailable@example.invalid>"
        body_tokens = ["Decision: NO", "one reference", "requires two", "105 words", "NOT_VERIFIABLE", "blocked", "example.invalid"]
        result["body"] = all(token.lower() in body.lower() for token in body_tokens)
        if source_pdf is not None and len(attachments) == 1:
            attachment = attachments[0]
            payload = attachment.get_payload(decode=True) or b""
            result["attachment"] = (
                attachment.get_filename() == PDF_NAME
                and attachment.get_content_type() == "application/pdf"
                and hashlib.sha256(payload).hexdigest() == SOURCE_HASHES[PDF_NAME]
                and payload == source_pdf.read_bytes()
            )
    except Exception as exc:
        result["errors"].append(f"parse:{type(exc).__name__}:{exc}")  # type: ignore[union-attr]
    return result


def unfold_ics(text: str) -> list[str]:
    lines: list[str] = []
    for raw in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        if raw.startswith((" ", "\t")) and lines:
            lines[-1] += raw[1:]
        elif raw:
            lines.append(raw)
    return lines


def calendar_audit(path: Path) -> dict[str, object]:
    result: dict[str, object] = {"errors": [], "text": "", "event": False, "safe": False, "event_count": 0}
    try:
        raw = path.read_bytes()
        text = raw.decode("utf-8")
        lines = unfold_ics(text)
        result["text"] = " ".join(lines)
        if b"\r\n" not in raw or raw.replace(b"\r\n", b"").find(b"\n") >= 0:
            result["errors"].append("not_strict_crlf")  # type: ignore[union-attr]
        result["event_count"] = lines.count("BEGIN:VEVENT")
        properties: dict[str, list[str]] = {}
        in_event = False
        for line in lines:
            if line == "BEGIN:VEVENT":
                in_event = True
                continue
            if line == "END:VEVENT":
                in_event = False
                continue
            if in_event and ":" in line:
                key, value = line.split(":", 1)
                properties.setdefault(key.split(";", 1)[0].upper(), []).append(value)
        description = " ".join(properties.get("DESCRIPTION", [])).replace("\\;", ";").replace("\\,", ",")
        result["event"] = (
            result["event_count"] == 1
            and properties.get("UID") == ["officebench-3-19-20-fellowship-followup@evocowork.invalid"]
            and properties.get("DTSTART") == ["20240531T130000Z"]
            and properties.get("SUMMARY") == ["Fellowship compliance follow-up with officer"]
            and properties.get("STATUS") == ["TENTATIVE"]
            and properties.get("X-EVOCOWORK-DECISION") == ["NO"]
            and "only one of two required references" in description
            and "NOT_VERIFIABLE" in description
        )
        forbidden = [key for key in ("DTEND", "DURATION", "LOCATION", "ATTENDEE", "ORGANIZER") if key in properties]
        result["safe"] = not forbidden
        if forbidden:
            result["errors"].append(f"forbidden:{forbidden}")  # type: ignore[union-attr]
    except Exception as exc:
        result["errors"].append(f"parse:{type(exc).__name__}:{exc}")  # type: ignore[union-attr]
    return result


def render_office(artifact_dir: Path) -> dict[str, object]:
    result: dict[str, object] = {"xlsx_pages": 0, "docx_pages": 0, "errors": [], "logs": []}
    with tempfile.TemporaryDirectory(prefix="fellowship-judge-render-") as temporary:
        base = Path(temporary)
        output = base / "output"
        profile = base / "profile"
        output.mkdir()
        for filename, key in ((XLSX_NAME, "xlsx_pages"), (DOCX_NAME, "docx_pages")):
            path = artifact_dir / filename
            if not path.is_file():
                result["errors"].append(f"missing:{filename}")  # type: ignore[union-attr]
                continue
            try:
                process = subprocess.run(
                    [
                        "libreoffice", "--headless", f"-env:UserInstallation=file://{profile}-{key}",
                        "--convert-to", "pdf", "--outdir", str(output), str(path),
                    ],
                    stdout=subprocess.PIPE,
                    stderr=subprocess.STDOUT,
                    text=True,
                    timeout=300,
                )
                result["logs"].append(process.stdout[-1000:])  # type: ignore[union-attr]
                pdf = output / f"{path.stem}.pdf"
                if process.returncode != 0 or not pdf.is_file():
                    result["errors"].append(f"render:{filename}:{process.returncode}")  # type: ignore[union-attr]
                    continue
                info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=60)
                match = re.search(r"^Pages:\s+(\d+)", info.stdout, flags=re.MULTILINE)
                if info.returncode != 0 or not match:
                    result["errors"].append(f"pdfinfo:{filename}")  # type: ignore[union-attr]
                    continue
                result[key] = int(match.group(1))
            except Exception as exc:
                result["errors"].append(f"render:{filename}:{type(exc).__name__}:{exc}")  # type: ignore[union-attr]
    return result


def unsupported_claims(texts: list[str], email_to: str) -> list[str]:
    joined = " ".join(texts)
    patterns = {
        "invented_second_reference": r"Dr\.?\s+(?:Jane|Mary|Robert|Emily)\s+[A-Z][a-z]+",
        "invented_deadline": r"(?:deadline|due date)\s+(?:is|was|:)\s*(?:20\d{2}|May|June|July|\d{1,2}/\d{1,2})",
        "invented_submission": r"(?:was|has been)\s+submitted\s+(?:on time|before|after|late)",
        "invented_eligibility": r"(?:is|was)\s+(?:eligible|approved|awarded|selected for the fellowship)",
        "invented_real_officer_address": r"officer@[A-Za-z0-9.-]+\.(?:com|org|edu|net)",
    }
    found = [name for name, pattern in patterns.items() if re.search(pattern, joined, flags=re.IGNORECASE)]
    if email_to != "Fellowship Officer <address-unavailable@example.invalid>":
        found.append("unexpected_recipient")
    return found


def criterion(identifier: str, passed: bool, evidence: object) -> dict[str, object]:
    description, weight, kind, cap = SPECS[identifier]
    raw: dict[str, object] = {"description": description, "weight": weight, "type": kind, "passed": passed}
    if cap is not None:
        raw["score_cap_on_failure"] = cap
    return {"criterion_id": identifier, "score": 1.0 if passed else 0.0, "evidence": evidence, "method": "deterministic", "raw": raw}


def evaluate(artifact_dir: Path) -> dict[str, object]:
    task_root = Path(__file__).resolve().parents[1]
    truth, source_errors, source_paths = validate_sources(task_root)
    delivered = sorted(path.name for path in artifact_dir.iterdir()) if artifact_dir.is_dir() else []
    exact_files = delivered == sorted(REQUIRED_FILES)
    packages = {
        name: inspect_ooxml(artifact_dir / name)
        for name in (XLSX_NAME, DOCX_NAME)
        if (artifact_dir / name).is_file()
    }
    package_safe = len(packages) == 2 and all(
        not package["broken"] and not package["macros"] and not package["external"] and not package["missing_targets"]
        for package in packages.values()
    )
    workbook = workbook_audit(artifact_dir / XLSX_NAME, truth) if (artifact_dir / XLSX_NAME).is_file() and not packages.get(XLSX_NAME, {}).get("broken") else {"errors": ["unavailable"], "policy": False, "application": False, "formulas": False, "decision": False, "boundary": False, "delivery": False, "native": {}, "text": ""}
    document = document_audit(artifact_dir / DOCX_NAME) if (artifact_dir / DOCX_NAME).is_file() and not packages.get(DOCX_NAME, {}).get("broken") else {"errors": ["unavailable"], "text": "", "headings": [], "tables": 0, "structure": False, "truth": False, "boundary": False}
    email = email_audit(artifact_dir / EML_NAME, source_paths.get(PDF_NAME)) if (artifact_dir / EML_NAME).is_file() else {"errors": ["unavailable"], "text": "", "headers": False, "body": False, "attachment": False, "safe_recipient": False}
    calendar = calendar_audit(artifact_dir / ICS_NAME) if (artifact_dir / ICS_NAME).is_file() else {"errors": ["unavailable"], "text": "", "event": False, "safe": False, "event_count": 0}
    render = render_office(artifact_dir) if (artifact_dir / XLSX_NAME).is_file() or (artifact_dir / DOCX_NAME).is_file() else {"xlsx_pages": 0, "docx_pages": 0, "errors": ["unavailable"], "logs": []}
    native = workbook.get("native", {})
    native_ok = (
        native.get("tables") == 4
        and native.get("charts") == 1
        and native.get("freeze") is True
        and int(native.get("conditional_format_ranges", 0)) >= 2
        and int(native.get("conditional_format_rules", 0)) >= 5
    )
    render_ok = not render["errors"] and render["xlsx_pages"] == 4 and render["docx_pages"] == 2
    email_to = ""
    if (artifact_dir / EML_NAME).is_file():
        try:
            email_to = normalize(BytesParser(policy=policy.default).parse((artifact_dir / EML_NAME).open("rb")).get("To"))
        except Exception:
            pass
    unsupported = unsupported_claims(
        [normalize(workbook.get("text")), normalize(document.get("text")), normalize(email.get("text")), normalize(calendar.get("text"))],
        email_to,
    )
    cross_tokens = ["NO", "Alice Johnson", "105", "one", "two", "NOT_VERIFIABLE"]
    cross_texts = [normalize(workbook.get("text")), normalize(document.get("text")), normalize(email.get("text")), normalize(calendar.get("text"))]
    cross_ok = all(all(token.lower() in text.lower() for token in cross_tokens) for text in cross_texts)
    suspicious = []
    for filename in REQUIRED_FILES:
        path = artifact_dir / filename
        if path.is_file():
            raw = path.read_bytes().lower()
            suspicious.extend(f"{filename}:{marker}" for marker in SUSPICIOUS_MARKERS if marker.encode("utf-8") in raw)

    booleans = {
        "R001": exact_files and not source_errors and package_safe,
        "R002": bool(workbook.get("policy")),
        "R003": bool(workbook.get("application")),
        "R004": bool(workbook.get("formulas")),
        "R005": bool(workbook.get("decision")) and bool(document.get("truth")),
        "R006": bool(workbook.get("boundary")) and bool(document.get("boundary")),
        "R007": native_ok,
        "R008": bool(document.get("structure")),
        "R009": bool(document.get("truth")) and bool(document.get("boundary")),
        "R010": bool(email.get("headers")) and bool(email.get("body")) and bool(email.get("attachment")) and bool(email.get("safe_recipient")),
        "R011": bool(calendar.get("event")) and bool(calendar.get("safe")),
        "R012": cross_ok and not unsupported,
        "R013": render_ok,
    }
    booleans.update({
        "R014": booleans["R002"] and booleans["R003"] and booleans["R004"] and booleans["R005"],
        "R015": not unsupported,
        "R016": exact_files and package_safe and booleans["R004"] and native_ok and not suspicious,
        "R017": booleans["R010"],
        "R018": booleans["R011"],
    })
    evidence = {
        "R001": {"source_errors": source_errors, "delivered": delivered, "packages": packages},
        "R002": {"policy": workbook.get("policy"), "errors": workbook.get("errors")},
        "R003": {"application": workbook.get("application"), "truth": truth},
        "R004": {"formulas": workbook.get("formulas")},
        "R005": {"workbook": workbook.get("decision"), "word": document.get("truth")},
        "R006": {"workbook": workbook.get("boundary"), "word": document.get("boundary")},
        "R007": native,
        "R008": {"tables": document.get("tables"), "headings": document.get("headings"), "errors": document.get("errors")},
        "R009": {"truth": document.get("truth"), "boundary": document.get("boundary")},
        "R010": email,
        "R011": calendar,
        "R012": {"cross": cross_ok, "unsupported": unsupported},
        "R013": {"render": render, "native": native},
        "R014": {"policy": workbook.get("policy"), "application": workbook.get("application"), "formulas": workbook.get("formulas"), "decision": workbook.get("decision")},
        "R015": {"unsupported": unsupported},
        "R016": {"files": exact_files, "packages": package_safe, "formulas": workbook.get("formulas"), "native": native_ok, "suspicious": suspicious},
        "R017": email,
        "R018": calendar,
    }
    results = [criterion(identifier, booleans[identifier], evidence[identifier]) for identifier in SPECS]
    positive_ids = [identifier for identifier, spec in SPECS.items() if spec[1] > 0]
    denominator = sum(SPECS[identifier][1] for identifier in positive_ids)
    reward = sum(SPECS[identifier][1] for identifier in positive_ids if booleans[identifier]) / denominator
    hurdle_failures = []
    for identifier, (_, weight, kind, cap) in SPECS.items():
        if kind == "hurdle" and not booleans[identifier]:
            hurdle_failures.append(identifier)
            reward = min(reward, float(cap))
        elif kind == "penalty" and not booleans[identifier]:
            reward += weight / 100.0
    reward = max(0.0, min(1.0, reward))
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [sha256(artifact_dir / name) for name in REQUIRED_FILES if (artifact_dir / name).is_file()],
        "criterion_results": results,
        "reward": reward,
        "passed": reward >= 0.8 and not hurdle_failures,
        "hurdle_failures": hurdle_failures,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    result = evaluate(args.artifact_dir)
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{result['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
