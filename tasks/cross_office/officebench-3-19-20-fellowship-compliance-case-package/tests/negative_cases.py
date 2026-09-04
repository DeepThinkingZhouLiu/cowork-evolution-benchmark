#!/usr/bin/env python3
"""奖学金合规案件包的 v2 分层负例。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import shutil
import sys
import tempfile
from email import policy
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.parser import BytesParser
from pathlib import Path
from typing import Callable

from docx import Document
from openpyxl import Workbook, load_workbook


ROOT=Path(__file__).resolve().parents[1]
GOLDEN=ROOT/"solution"/"Golden"
INPUTS=ROOT/"environment"/"input_files"
REQUIRED=("Fellowship_Compliance_Audit.xlsx","Fellowship_Compliance_Decision.docx","Fellowship_Application_Forward.eml","Fellowship_Compliance_Followup.ics")
Mutator=Callable[[Path],None]


def _load_judge():
    path=Path(__file__).with_name("judge.py");spec=importlib.util.spec_from_file_location("fellowship_v2_judge_for_negatives",path)
    if spec is None or spec.loader is None:raise RuntimeError("无法载入 v2 Judge")
    module=importlib.util.module_from_spec(spec);sys.modules[spec.name]=module;spec.loader.exec_module(module);return module


JUDGE=_load_judge()


def _copy(target:Path)->None:
    target.mkdir(parents=True,exist_ok=True)
    for name in REQUIRED:shutil.copy2(GOLDEN/name,target/name)


def _excel(sheet:str,cell:str,value:object)->Mutator:
    def apply(root:Path)->None:
        path=root/REQUIRED[0];book=load_workbook(path);book[sheet][cell]=value;book.save(path);book.close()
    return apply


def _excel_formula(sheet:str,cell:str,value:object)->Mutator:
    return _excel(sheet,cell,value)


def _excel_remove_chart(root:Path)->None:
    path=root/REQUIRED[0];book=load_workbook(path);book["Decision_Summary"]._charts=[];book.save(path);book.close()


def _excel_remove_freeze(root:Path)->None:
    path=root/REQUIRED[0];book=load_workbook(path);book["Policy_Criteria"].freeze_panes=None;book.save(path);book.close()


def _excel_remove_cf(root:Path)->None:
    path=root/REQUIRED[0];book=load_workbook(path);book["Policy_Criteria"].conditional_formatting._cf_rules.clear();book.save(path);book.close()


def _word_replace(old:str,new:str)->Mutator:
    def apply(root:Path)->None:
        path=root/REQUIRED[1];document=Document(path);locations=list(document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:locations.extend(cell.paragraphs)
        for paragraph in locations:
            if old in paragraph.text:
                paragraph.text=paragraph.text.replace(old,new,1);document.save(path);return
        raise AssertionError(f"Word 未找到 {old}")
    return apply


def _word_unsupported(root:Path)->None:
    path=root/REQUIRED[1];document=Document(path);document.add_paragraph("The application was submitted on time before the deadline and is eligible for the fellowship.");document.save(path)


def _word_table_cell(table_index:int,row_index:int,column_index:int,value:str)->Mutator:
    def apply(root:Path)->None:
        path=root/REQUIRED[1];document=Document(path);document.tables[table_index].rows[row_index].cells[column_index].text=value;document.save(path)
    return apply


def _save_message(message,path:Path)->None:
    with path.open("wb") as handle:BytesGenerator(handle,policy=policy.SMTP,mangle_from_=False).flatten(message)


def _email_header(header:str,value:str)->Mutator:
    def apply(root:Path)->None:
        path=root/REQUIRED[2];message=BytesParser(policy=policy.default).parse(path.open("rb"));message.replace_header(header,value);_save_message(message,path)
    return apply


def _email_body_replace(old:str,new:str)->Mutator:
    def apply(root:Path)->None:
        path=root/REQUIRED[2];message=BytesParser(policy=policy.default).parse(path.open("rb"));body=message.get_body(preferencelist=("plain",));text=body.get_content();assert old in text;body.set_content(text.replace(old,new,1),subtype="plain",charset="utf-8");_save_message(message,path)
    return apply


def _email_drop_attachment(root:Path)->None:
    path=root/REQUIRED[2];message=BytesParser(policy=policy.default).parse(path.open("rb"));del message.get_payload()[1:];_save_message(message,path)


def _email_modify_attachment(root:Path)->None:
    path=root/REQUIRED[2];message=BytesParser(policy=policy.default).parse(path.open("rb"));attachment=next(message.iter_attachments());attachment.set_content(b"changed-pdf",maintype="application",subtype="pdf",cte="base64");attachment.add_header("Content-Disposition","attachment",filename="fellowship_application.pdf");_save_message(message,path)


def _ics_replace(old:bytes,new:bytes)->Mutator:
    def apply(root:Path)->None:
        path=root/REQUIRED[3];raw=path.read_bytes();assert old in raw;path.write_bytes(raw.replace(old,new,1))
    return apply


def _moderate_communication_domain(root:Path)->None:
    email_path=root/REQUIRED[2];message=BytesParser(policy=policy.default).parse(email_path.open("rb"));message.replace_header("To","Fellowship Officer <officer@school.edu>");message.replace_header("Subject","Fellowship award notice");message.replace_header("X-EvoCowork-Delivery-Status","READY");message.replace_header("X-EvoCowork-Decision","YES");body=message.get_body(preferencelist=("plain",));body.set_content("Fellowship Officer,\n\nDecision: YES. The application was submitted on time and is ready for award.\n",subtype="plain",charset="utf-8");del message.get_payload()[1:];_save_message(message,email_path)
    (root/REQUIRED[3]).write_bytes(
        b"BEGIN:VCALENDAR\r\nVERSION:2.0\r\nPRODID:-//EvoCowork//Fellowship//EN\r\nBEGIN:VEVENT\r\n"
        b"UID:wrong@school.edu\r\nDTSTART:20240531T150000Z\r\nDTEND:20240531T160000Z\r\n"
        b"SUMMARY:Fellowship award meeting\r\nSTATUS:CONFIRMED\r\nLOCATION:Room 7\r\n"
        b"ATTENDEE:mailto:officer@school.edu\r\nDESCRIPTION:Decision YES; application complete.\r\n"
        b"X-EVOCOWORK-DECISION:YES\r\nEND:VEVENT\r\nEND:VCALENDAR\r\n"
    )


def _severe_valid_incomplete(root:Path)->None:
    xlsx=root/REQUIRED[0];book=Workbook();book.remove(book.active)
    for name,headers in (("Policy_Criteria",JUDGE.V1.POLICY_HEADERS),("Application_Evidence",JUDGE.V1.APPLICATION_HEADERS),("Decision_Summary",JUDGE.V1.SUMMARY_HEADERS),("Delivery_Control",JUDGE.V1.DELIVERY_HEADERS)):
        sheet=book.create_sheet(name);sheet.append(headers)
    book.save(xlsx);book.close()
    document=Document();document.add_paragraph("Fellowship Application Compliance Decision");document.add_paragraph("The detailed review is incomplete.");document.save(root/REQUIRED[1])
    message=EmailMessage();message["To"]="Fellowship Officer <address-unavailable@example.invalid>";message["Subject"]="Fellowship application compliance review - Alice Johnson";message.set_content("The detailed fellowship review is incomplete.");_save_message(message,root/REQUIRED[2])
    (root/REQUIRED[3]).write_bytes(b"BEGIN:VCALENDAR\r\nVERSION:2.0\r\nPRODID:-//EvoCowork//Incomplete//EN\r\nEND:VCALENDAR\r\n")


def _input_only(root:Path)->None:
    for path in list(root.iterdir()):
        if path.is_file():path.unlink()
    for source in INPUTS.iterdir():
        if source.is_file():shutil.copy2(source,root/source.name)


PARTIALS:dict[str,tuple[Mutator,bool]]={
    "policy_one_id_wrong":(_excel("Policy_Criteria","A2","P99"),False),
    "policy_one_text_wrong":(_excel("Policy_Criteria","B2","The applicant name is optional."),False),
    "policy_one_evidence_wrong":(_excel("Policy_Criteria","C4","Statement contains 99 words."),False),
    "policy_one_status_wrong":(_excel("Policy_Criteria","D5","PASS"),False),
    "policy_one_determinative_wrong":(_excel("Policy_Criteria","E5","No"),False),
    "policy_one_note_wrong":(_excel("Policy_Criteria","F6","Submitted late."),False),
    "application_name_wrong":(_excel("Application_Evidence","B2","Alicia Johnson"),False),
    "application_word_count_wrong":(_excel("Application_Evidence","B5",99),False),
    "application_reference_count_wrong":(_excel("Application_Evidence","B6",2),False),
    "application_deadline_invented":(_excel("Application_Evidence","B8","2024-05-15"),False),
    "formula_one_status_flattened":(_excel_formula("Policy_Criteria","D2","PASS"),True),
    "formula_one_status_wrong_reference":(_excel_formula("Policy_Criteria","D5",'=IF(Application_Evidence!B5>=2,"PASS","FAIL")'),True),
    "formula_overall_decision_flattened":(_excel_formula("Decision_Summary","B5","NO"),True),
    "delivery_email_status_wrong":(_excel("Delivery_Control","D4","READY"),False),
    "native_chart_removed":(_excel_remove_chart,False),
    "native_freeze_removed":(_excel_remove_freeze,False),
    "native_conditional_format_removed":(_excel_remove_cf,False),
    "word_heading_missing":(_word_replace("Policy-by-Policy Assessment","Policy Review"),False),
    "word_decision_conflict":(_word_replace("Decision: NO.","Decision: YES."),False),
    "word_applicant_wrong":(_word_table_cell(1,1,1,"Alicia Johnson"),False),
    "word_one_policy_status_wrong":(_word_table_cell(2,5,2,"PASS"),False),
    "word_one_remediation_missing":(_word_replace("Obtain one additional reference","No remediation required"),False),
    "word_time_boundary_missing":(_word_replace("not treated as proof of on-time or late submission","treated as submitted on time"),False),
    "word_unsupported_eligibility":(_word_unsupported,False),
    "email_wrong_subject":(_email_header("Subject","Fellowship application update"),False),
    "email_wrong_recipient":(_email_header("To","Fellowship Officer <officer@school.edu>"),False),
    "email_wrong_delivery_status":(_email_header("X-EvoCowork-Delivery-Status","READY"),False),
    "email_decision_conflict":(_email_body_replace("Decision: NO.","Decision: YES."),False),
    "email_drop_attachment":(_email_drop_attachment,False),
    "email_modify_attachment":(_email_modify_attachment,False),
    "ics_wrong_time":(_ics_replace(b"DTSTART:20240531T130000Z",b"DTSTART:20240531T140000Z"),False),
    "ics_wrong_status":(_ics_replace(b"STATUS:TENTATIVE",b"STATUS:CONFIRMED"),False),
    "ics_wrong_summary":(_ics_replace(b"SUMMARY:Fellowship compliance follow-up with officer",b"SUMMARY:Fellowship award meeting"),False),
    "ics_add_dtend":(_ics_replace(b"DTSTART:20240531T130000Z\r\n",b"DTSTART:20240531T130000Z\r\nDTEND:20240531T140000Z\r\n"),False),
    "ics_add_location":(_ics_replace(b"STATUS:TENTATIVE\r\n",b"STATUS:TENTATIVE\r\nLOCATION:Room 7\r\n"),False),
    "moderate_communication_domain_conflict":(_moderate_communication_domain,False),
    "severe_valid_but_incomplete_package":(_severe_valid_incomplete,False),
    "unexpected_debug_file":(lambda root:(root/"debug.txt").write_text("temporary\n",encoding="utf-8"),False),
}


def main()->None:
    parser=argparse.ArgumentParser();parser.add_argument("--report",type=Path);args=parser.parse_args();golden_recalc,_=JUDGE._recalculate(GOLDEN/REQUIRED[0]);golden_render=JUDGE._render(GOLDEN);original_recalc,original_render=JUDGE._recalculate,JUDGE._render
    assert JUDGE.evaluate(GOLDEN)["reward"]==1.0
    def fast_recalc(path:Path):
        book=load_workbook(path,data_only=False);values={}
        for sheet in book.worksheets:
            rows=[]
            for row_index,row in enumerate(sheet.iter_rows(),1):
                current=[]
                for column_index,cell in enumerate(row,1):
                    value=cell.value
                    if isinstance(value,str) and value.startswith("="):value=JUDGE._value(golden_recalc,sheet.title,row_index,column_index)
                    current.append(value)
                rows.append(current)
            values[sheet.title]=rows
        book.close();return values,"task-specific cached Golden formula values"
    def evaluate_case(case:Path,real_recalc:bool)->float:
        JUDGE._recalculate=original_recalc if real_recalc else fast_recalc;JUDGE._render=lambda _artifact:golden_render
        try:return float(JUDGE.evaluate(case)["reward"])
        finally:JUDGE._recalculate,JUDGE._render=original_recalc,original_render
    records=[]
    with tempfile.TemporaryDirectory(prefix="fellowship-v2-negatives-") as raw:
        base=Path(raw)
        for name in REQUIRED:
            for mode in ("missing","empty","corrupt"):
                case=base/f"fatal-{mode}-{name}";_copy(case);target=case/name
                if mode=="missing":target.unlink()
                elif mode=="empty":target.write_bytes(b"")
                else:target.write_bytes(b"not-a-valid-task-artifact")
                reward=evaluate_case(case,False);assert reward==0,(mode,name,reward);records.append({"case":f"{mode}:{name}","layer":"fatal_zero","reward":reward})
        empty=base/"fatal-empty";empty.mkdir();reward=evaluate_case(empty,False);assert reward==0;records.append({"case":"empty-output","layer":"fatal_zero","reward":reward})
        input_only=base/"fatal-input-only";_copy(input_only);_input_only(input_only);reward=evaluate_case(input_only,False);assert reward==0;records.append({"case":"unmodified-input-only","layer":"fatal_zero","reward":reward})
        for name,(mutator,real_recalc) in PARTIALS.items():
            case=base/f"partial-{name}";_copy(case);mutator(case);reward=evaluate_case(case,real_recalc);assert 0<reward<1,(name,reward);records.append({"case":name,"layer":"partial","reward":reward})
    partial=[float(record["reward"]) for record in records if record["layer"]=="partial"];report={"total":len(records),"fatal_zero":sum(record["layer"]=="fatal_zero" for record in records),"partial_continuous":len(partial),"partial_reward_range":[min(partial),max(partial)],"cases":records}
    if args.report:args.report.parent.mkdir(parents=True,exist_ok=True);args.report.write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=="__main__":main()
