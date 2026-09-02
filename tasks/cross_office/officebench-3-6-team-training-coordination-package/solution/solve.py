#!/usr/bin/env python3
"""OfficeBench 3-6 团队培训协调包的逐题参考解法。"""

from __future__ import annotations

import argparse
import hashlib
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import date, time
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.policy import SMTP
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


SOURCE_NAME = "team_training_schedule.xlsx"
SOURCE_SHA256 = "a5628f45d9f48eccd8424634fda4d3c9048200571cc03db1e6d34cf9447498f4"
XLSX_NAME = "Team_Training_Coordination_Audit.xlsx"
DOCX_NAME = "Team_Training_Coordination_Brief.docx"
ICS_NAME = "Team_Training_Master.ics"
MBOX_NAME = "Team_Training_Reminders.mbox"
TRAINING_DATE = date(2020, 5, 8)
FIXED_ZIP_TIME = (2020, 5, 1, 10, 0, 0)
FIXED_CORE_TIME = "2020-05-01T10:00:00Z"

NAVY = "17365D"
BLUE = "4472C4"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EAF2F8"
GREEN = "E2F0D9"
AMBER = "FFF2CC"
RED = "F4CCCC"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
THIN_GRAY = Side(style="thin", color="B7C9D6")


@dataclass(frozen=True)
class Member:
    source_row: int
    name: str
    section: str
    source_time: str
    start: time
    end: time


EXPECTED_ROWS = (
    ("Alice", "Section A", "9:00 AM - 11:00 AM"),
    ("Bob", "Section B", "11:00 AM - 1:00 PM"),
    ("Carol", "Section C", "2:00 PM - 4:00 PM"),
    ("David", "Section A", "9:00 AM - 11:00 AM"),
    ("Eva", "Section B", "11:00 AM - 1:00 PM"),
    ("Frank", "Section C", "2:00 PM - 4:00 PM"),
    ("Grace", "Section A", "9:00 AM - 11:00 AM"),
    ("Hank", "Section B", "11:00 AM - 1:00 PM"),
    ("Ivy", "Section C", "2:00 PM - 4:00 PM"),
    ("Jack", "Section A", "9:00 AM - 11:00 AM"),
    ("Kelly", "Section B", "11:00 AM - 1:00 PM"),
    ("Leo", "Section C", "2:00 PM - 4:00 PM"),
)

SESSION_TIMES = {
    "Section A": (time(9, 0), time(11, 0)),
    "Section B": (time(11, 0), time(13, 0)),
    "Section C": (time(14, 0), time(16, 0)),
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_clock(value: str) -> time:
    match = re.fullmatch(r"(\d{1,2}):(\d{2}) (AM|PM)", value.strip())
    if not match:
        raise ValueError(f"无法解析时间：{value}")
    hour = int(match.group(1)) % 12
    if match.group(3) == "PM":
        hour += 12
    return time(hour, int(match.group(2)))


def parse_source(input_dir: Path) -> list[Member]:
    source = input_dir / SOURCE_NAME
    if not source.is_file() or sha256(source) != SOURCE_SHA256:
        raise ValueError("培训源表缺失或 SHA256 与冻结真值不一致")
    workbook = load_workbook(source, data_only=False, read_only=True)
    if workbook.sheetnames != ["Training Schedule"]:
        raise ValueError("源工作簿应且仅应包含 Training Schedule")
    sheet = workbook["Training Schedule"]
    headers = tuple(sheet.cell(1, column).value for column in range(1, 4))
    if headers != ("Name", "Section", "Time") or sheet.max_row != 13 or sheet.max_column != 3:
        raise ValueError("源表结构与上游冻结版本不一致")
    members: list[Member] = []
    actual_rows: list[tuple[str, str, str]] = []
    for row in range(2, 14):
        name, section, source_time = (sheet.cell(row, column).value for column in range(1, 4))
        actual = (str(name), str(section), str(source_time))
        actual_rows.append(actual)
        pieces = actual[2].split(" - ")
        if len(pieces) != 2:
            raise ValueError(f"第 {row} 行时间范围无法解析")
        start, end = parse_clock(pieces[0]), parse_clock(pieces[1])
        if actual[1] not in SESSION_TIMES or (start, end) != SESSION_TIMES[actual[1]]:
            raise ValueError(f"第 {row} 行 Section 与时间不匹配")
        members.append(Member(row, actual[0], actual[1], actual[2], start, end))
    if tuple(actual_rows) != EXPECTED_ROWS:
        raise ValueError("源表的 12 行人员真值已变化")
    return members


def section_slug(section: str) -> str:
    return section.lower().replace(" ", "-")


def section_uid(section: str) -> str:
    return f"{section_slug(section)}-training-20200508@evocowork.invalid"


def members_by_section(members: list[Member]) -> dict[str, list[Member]]:
    grouped = {section: [] for section in SESSION_TIMES}
    for member in members:
        grouped[member.section].append(member)
    return grouped


def normalize_ooxml(path: Path) -> None:
    """统一 OOXML ZIP 顺序、时间戳、核心属性与 LibreOffice 随机图表轴 ID。"""

    with tempfile.TemporaryDirectory(prefix="training-ooxml-") as temporary:
        target = Path(temporary) / path.name
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
            target, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
        ) as destination:
            for name in sorted(source.namelist()):
                data = source.read(name)
                if name == "docProps/core.xml":
                    try:
                        root = etree.fromstring(data)
                        namespaces = {"dcterms": "http://purl.org/dc/terms/"}
                        for tag in ("created", "modified"):
                            for node in root.xpath(f"//dcterms:{tag}", namespaces=namespaces):
                                node.text = FIXED_CORE_TIME
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except Exception:
                        pass
                if name.startswith("xl/charts/chart") and name.endswith(".xml"):
                    try:
                        root = etree.fromstring(data)
                        namespace = "http://schemas.openxmlformats.org/drawingml/2006/chart"
                        stable_ids: dict[str, str] = {}
                        for node in root.xpath("//c:axId", namespaces={"c": namespace}):
                            old_id = node.get("val")
                            if old_id is not None and old_id not in stable_ids:
                                stable_ids[old_id] = str(20_000_001 + len(stable_ids))
                        for node in root.xpath("//c:axId | //c:crossAx", namespaces={"c": namespace}):
                            old_id = node.get("val")
                            if old_id in stable_ids:
                                node.set("val", stable_ids[old_id])
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    except Exception:
                        pass
                info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                destination.writestr(info, data)
        shutil.copy2(target, path)


def recalculate_workbook(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="training-recalc-") as temporary:
        output_dir = Path(temporary) / "output"
        profile_dir = Path(temporary) / "profile"
        output_dir.mkdir()
        process = subprocess.run(
            [
                "libreoffice",
                "--headless",
                f"-env:UserInstallation=file://{profile_dir}",
                "--convert-to",
                "xlsx",
                "--outdir",
                str(output_dir),
                str(path),
            ],
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            timeout=300,
        )
        converted = output_dir / path.name
        if process.returncode != 0 or not converted.is_file():
            raise RuntimeError(f"LibreOffice 公式重算失败：{process.stdout[-1000:]}")
        shutil.copy2(converted, path)


def restore_frozen_rows(path: Path, sheet_count: int) -> None:
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with tempfile.TemporaryDirectory(prefix="training-freeze-") as temporary:
        target = Path(temporary) / path.name
        sheet_names = {f"xl/worksheets/sheet{index}.xml" for index in range(1, sheet_count + 1)}
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w") as destination:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename in sheet_names:
                    root = etree.fromstring(data)
                    sheet_view = root.find(f"{{{namespace}}}sheetViews/{{{namespace}}}sheetView")
                    if sheet_view is not None:
                        for existing in sheet_view.findall(f"{{{namespace}}}pane"):
                            sheet_view.remove(existing)
                        pane = etree.Element(f"{{{namespace}}}pane")
                        pane.set("ySplit", "1")
                        pane.set("topLeftCell", "A2")
                        pane.set("activePane", "bottomLeft")
                        pane.set("state", "frozen")
                        sheet_view.insert(0, pane)
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                destination.writestr(info, data)
        shutil.copy2(target, path)


def add_table(sheet, name: str, reference: str) -> None:
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False
    )
    sheet.add_table(table)


def style_sheet(sheet, widths: dict[str, float], landscape: bool = True) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.oddHeader.center.text = "&BTeam Training Coordination Audit&B"
    sheet.oddFooter.right.text = "Page &P of &N"
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.45
    sheet.page_margins.bottom = 0.45
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=BLUE))
    sheet.row_dimensions[1].height = 34
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=9, color="1F1F1F")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(left=THIN_GRAY, right=THIN_GRAY, bottom=THIN_GRAY)


def build_workbook(members: list[Member], output_path: Path) -> None:
    grouped = members_by_section(members)
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.creator = "EvoCowork benchmark oracle"
    workbook.title = "Team Training Coordination Audit"
    workbook.subject = "OfficeBench 3-6 source-backed coordination controls"
    workbook.description = "Formula-backed audit of the frozen OfficeBench 3-6 training schedule."
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"

    roster = workbook.create_sheet("Roster_Control")
    roster.append(
        [
            "Source Row",
            "Name",
            "Section",
            "Source Time",
            "Start",
            "End",
            "Duration (Hours)",
            "Training Date",
            "Reserved Address",
            "Delivery Status",
            "Session UID",
        ]
    )
    for index, member in enumerate(members, start=2):
        roster.append(
            [
                member.source_row,
                member.name,
                member.section,
                member.source_time,
                member.start,
                member.end,
                f"=ROUND(MOD(F{index}-E{index},1)*24,2)",
                TRAINING_DATE,
                f'=LOWER(B{index})&"@example.invalid"',
                "BLOCKED_MISSING_REAL_EMAIL",
                f'=LOWER(SUBSTITUTE(C{index}," ","-"))&"-training-20200508@evocowork.invalid"',
            ]
        )
    style_sheet(
        roster,
        {"A": 14, "B": 16, "C": 14, "D": 25, "E": 12, "F": 12, "G": 18, "H": 16, "I": 31, "J": 32, "K": 49},
    )
    for row in range(2, 14):
        roster.row_dimensions[row].height = 31
        roster[f"E{row}"].number_format = "h:mm AM/PM"
        roster[f"F{row}"].number_format = "h:mm AM/PM"
        roster[f"H{row}"].number_format = "yyyy-mm-dd"
    roster.conditional_formatting.add(
        "J2:J13", FormulaRule(formula=['J2="BLOCKED_MISSING_REAL_EMAIL"'], fill=PatternFill("solid", fgColor=AMBER))
    )
    add_table(roster, "TrainingRosterControlTable", "A1:K13")
    roster.print_area = "A1:K13"

    sessions = workbook.create_sheet("Session_Control")
    sessions.append(
        [
            "Section",
            "Training Date",
            "Start",
            "End",
            "Duration (Hours)",
            "Member Count",
            "Members",
            "Previous End",
            "Gap (Minutes)",
            "Overlap Flag",
            "Session UID",
            "Calendar Status",
        ]
    )
    for row_index, (section, (start, end)) in enumerate(SESSION_TIMES.items(), start=2):
        names = " | ".join(member.name for member in grouped[section])
        previous_end = None if row_index == 2 else f"=D{row_index - 1}"
        gap_formula = None if row_index == 2 else f"=ROUND((C{row_index}-H{row_index})*1440,0)"
        overlap_formula = "N/A" if row_index == 2 else f'=IF(C{row_index}<H{row_index},"YES","NO")'
        sessions.append(
            [
                section,
                TRAINING_DATE,
                start,
                end,
                f"=ROUND(MOD(D{row_index}-C{row_index},1)*24,2)",
                f'=COUNTIF(\'Roster_Control\'!$C$2:$C$13,A{row_index})',
                names,
                previous_end,
                gap_formula,
                overlap_formula,
                section_uid(section),
                "TENTATIVE",
            ]
        )
    style_sheet(
        sessions,
        {"A": 15, "B": 16, "C": 12, "D": 12, "E": 18, "F": 17, "G": 42, "H": 15, "I": 15, "J": 15, "K": 50, "L": 18},
    )
    for row in range(2, 5):
        sessions.row_dimensions[row].height = 38
        sessions[f"B{row}"].number_format = "yyyy-mm-dd"
        for column in ("C", "D", "H"):
            sessions[f"{column}{row}"].number_format = "h:mm AM/PM"
    sessions.conditional_formatting.add(
        "J2:J4", FormulaRule(formula=['J2="YES"'], fill=PatternFill("solid", fgColor=RED))
    )
    sessions.conditional_formatting.add(
        "L2:L4", FormulaRule(formula=['L2="TENTATIVE"'], fill=PatternFill("solid", fgColor=AMBER))
    )
    add_table(sessions, "TrainingSessionControlTable", "A1:L4")
    sessions.print_area = "A1:L4"

    reminders = workbook.create_sheet("Reminder_Control")
    reminders.append(
        [
            "Source Row",
            "Name",
            "Section",
            "Reserved Recipient",
            "Subject",
            "Message ID",
            "Calendar Attachment",
            "Session UID",
            "Delivery Status",
        ]
    )
    for row_index, member in enumerate(members, start=2):
        slug = section_slug(member.section)
        reminders.append(
            [
                member.source_row,
                member.name,
                member.section,
                f'=LOWER(B{row_index})&"@example.invalid"',
                f'="Team training reminder - "&C{row_index}&" - 2020-05-08"',
                f"<{member.name.lower()}-training-20200508@evocowork.invalid>",
                f"{slug}-training.ics",
                section_uid(member.section),
                "BLOCKED_MISSING_REAL_EMAIL",
            ]
        )
    style_sheet(
        reminders,
        {"A": 12, "B": 16, "C": 14, "D": 31, "E": 48, "F": 48, "G": 29, "H": 49, "I": 32},
    )
    for row in range(2, 14):
        reminders.row_dimensions[row].height = 38
    reminders.conditional_formatting.add(
        "I2:I13", FormulaRule(formula=['I2="BLOCKED_MISSING_REAL_EMAIL"'], fill=PatternFill("solid", fgColor=AMBER))
    )
    add_table(reminders, "TrainingReminderControlTable", "A1:I13")
    reminders.print_area = "A1:I13"

    summary = workbook.create_sheet("Delivery_Summary")
    summary.append(["Metric", "Formula Value", "Expected", "Control Status", "Control Note"])
    metrics = [
        ("Source Rows", "=COUNTA('Roster_Control'!B2:B13)", 12, "All frozen source rows retained"),
        ("Sections", "=COUNTA('Session_Control'!A2:A4)", 3, "One master VEVENT per section"),
        ("Reminder Messages", "=COUNTA('Reminder_Control'!B2:B13)", 12, "One multipart draft per source row"),
        ("Blocked Drafts", '=COUNTIF(\'Reminder_Control\'!I2:I13,"BLOCKED_MISSING_REAL_EMAIL")', 12, "No real addresses were supplied"),
        ("Calendar Events", "=COUNTA('Session_Control'!K2:K4)", 3, "Stable UID per section"),
        ("Scheduled Hours", "=SUM('Session_Control'!E2:E4)", 6, "Three sessions of two hours"),
        ("Inter-session Gap (Minutes)", "=SUM('Session_Control'!I2:I4)", 60, "Only the B-to-C gap contributes"),
        ("Overlap Count", '=COUNTIF(\'Session_Control\'!J2:J4,"YES")', 0, "Touching at 11:00 is not overlap"),
    ]
    for row_index, (metric, formula, expected, note) in enumerate(metrics, start=2):
        summary.append([metric, formula, expected, f'=IF(B{row_index}=C{row_index},"PASS","FAIL")', note])
    style_sheet(summary, {"A": 34, "B": 18, "C": 15, "D": 20, "E": 58, "F": 3, "G": 16, "H": 16, "I": 16, "J": 16, "K": 16, "L": 16})
    for row in range(2, 10):
        summary.row_dimensions[row].height = 29
    summary.conditional_formatting.add(
        "D2:D9", FormulaRule(formula=['D2="PASS"'], fill=PatternFill("solid", fgColor=GREEN))
    )
    summary.conditional_formatting.add(
        "D2:D9", FormulaRule(formula=['D2="FAIL"'], fill=PatternFill("solid", fgColor=RED))
    )
    add_table(summary, "TrainingDeliverySummaryTable", "A1:E9")
    chart = BarChart()
    chart.style = 10
    chart.title = "Members per training section"
    chart.y_axis.title = "Members"
    chart.x_axis.title = "Section"
    chart.height = 7.2
    chart.width = 13.0
    chart.add_data(Reference(sessions, min_col=6, min_row=1, max_row=4), titles_from_data=True)
    chart.set_categories(Reference(sessions, min_col=1, min_row=2, max_row=4))
    summary.add_chart(chart, "G2")
    summary.print_area = "A1:M9"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    recalculate_workbook(output_path)
    restore_frozen_rows(output_path, 4)
    normalize_ooxml(output_path)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_width(cell, centimeters: float) -> None:
    properties = cell._tc.get_or_add_tcPr()
    width = properties.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        properties.append(width)
    width.set(qn("w:w"), str(int(centimeters * 567)))
    width.set(qn("w:type"), "dxa")


def add_docx_table(
    document: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    font_size: float = 7.5,
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, width in enumerate(widths):
        table.columns[index].width = Cm(width)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, NAVY)
        set_cell_width(cell, widths[index])
        cell.width = Cm(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(font_size)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
            set_cell_width(cells[index], widths[index])
            cells[index].width = Cm(widths[index])
            if row_index % 2 == 0:
                set_cell_shading(cells[index], PALE_BLUE)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(font_size)
    table.style = "Table Grid"


def add_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    run.font.name = "Aptos Display"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(23, 54, 93)


def add_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.05
    for run in paragraph.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8.8)


def build_document(members: list[Member], output_path: Path) -> None:
    grouped = members_by_section(members)
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Cm(1.25)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)
    section.header_distance = Cm(0.55)
    section.footer_distance = Cm(0.55)
    document.core_properties.author = "EvoCowork benchmark oracle"
    document.core_properties.title = "Team Training Coordination Brief"
    document.core_properties.subject = "Source-backed OfficeBench 3-6 training coordination"

    header = section.header.paragraphs[0]
    header.text = "EVOCOWORK  |  TRAINING COORDINATION CONTROL"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(7.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(68, 114, 196)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Source: OfficeBench 3-6  |  Frozen date: 2020-05-01  |  Draft controls only")
    for run in footer.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(7)
        run.font.color.rgb = RGBColor(89, 89, 89)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(1)
    run = title.add_run("Team Training Coordination Brief")
    run.font.name = "Aptos Display"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(23, 54, 93)
    subtitle = document.add_paragraph("Training date 2020-05-08  |  12 members  |  3 sections  |  12 reminder drafts blocked")
    subtitle.paragraph_format.space_after = Pt(8)
    for run in subtitle.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor(68, 114, 196)

    add_heading(document, "Executive Control Summary")
    add_body(
        document,
        "The frozen source workbook contains 12 rows: four members each in Section A, Section B and Section C. "
        "The upstream task fixes the next-week training date as 2020-05-08. All 12 reminder messages are drafts, "
        "not sent communications, because no real recipient email address is present in the source.",
    )
    add_docx_table(
        document,
        ["Section", "Date", "Floating local time", "Members", "Calendar status"],
        [
            [section_name, "2020-05-08", f"{start.strftime('%H:%M')}-{end.strftime('%H:%M')}", " | ".join(m.name for m in grouped[section_name]), "TENTATIVE"]
            for section_name, (start, end) in SESSION_TIMES.items()
        ],
        [2.3, 2.5, 3.2, 7.3, 2.7],
        font_size=7.5,
    )
    add_body(
        document,
        "Section A and Section B touch at 11:00 and do not overlap. Section B ends at 13:00; Section C begins at 14:00, "
        "leaving a 60-minute inter-session gap. Each section lasts exactly two hours.",
    )

    add_heading(document, "Evidence Boundaries")
    add_docx_table(
        document,
        ["Missing source field", "Required handling", "Control outcome"],
        [
            ["Real recipient emails", "Use non-deliverable .invalid placeholders only", "12 drafts BLOCKED_MISSING_REAL_EMAIL"],
            ["Time zone", "Keep DTSTART/DTEND as floating local time", "No TZID and no UTC conversion"],
            ["Location", "State that location is not provided", "No LOCATION property"],
            ["Organizer/attendee mailboxes", "Do not fabricate identities or addresses", "No ORGANIZER or ATTENDEE property"],
        ],
        [4.2, 7.4, 6.4],
        font_size=7.2,
    )

    document.add_page_break()
    spacer = document.add_paragraph(" ")
    spacer.paragraph_format.space_after = Pt(8)
    add_heading(document, "Source Roster and Reminder Mapping")
    roster_rows = []
    for member in members:
        roster_rows.append(
            [
                str(member.source_row),
                member.name,
                member.section,
                member.source_time,
                f"{member.name.lower()}@example.invalid",
                "BLOCKED",
            ]
        )
    add_docx_table(
        document,
        ["Row", "Name", "Section", "Source time", "Reserved recipient", "Status"],
        roster_rows,
        [0.9, 1.8, 2.1, 3.6, 6.2, 1.7],
        font_size=6.6,
    )

    add_heading(document, "Delivery Controls")
    add_docx_table(
        document,
        ["Artifact", "Required content", "Status"],
        [
            [XLSX_NAME, "4 native audit tables, formulas and editable section chart", "READY"],
            [DOCX_NAME, "Schedule, roster, evidence gaps and controls", "READY"],
            [ICS_NAME, "3 tentative floating-time events with stable UIDs", "READY"],
            [MBOX_NAME, "12 multipart drafts with section calendar attachments", "BLOCKED"],
        ],
        [7.2, 7.8, 1.8],
        font_size=7.0,
    )
    add_body(
        document,
        "Release condition: replace every reserved .invalid recipient with an approved real address and separately confirm the time zone and location. "
        "Until then, no reminder may be represented as sent and all calendar events remain tentative.",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    normalize_ooxml(output_path)


def escape_ics(value: str) -> str:
    return value.replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")


def fold_ics_line(line: str, limit: int = 73) -> list[str]:
    chunks: list[str] = []
    remaining = line
    first = True
    while len(remaining.encode("utf-8")) > limit:
        cut = min(len(remaining), limit)
        while cut > 0 and len(remaining[:cut].encode("utf-8")) > limit:
            cut -= 1
        if cut == 0:
            raise ValueError("ICS 折行失败")
        chunks.append(("" if first else " ") + remaining[:cut])
        remaining = remaining[cut:]
        first = False
        limit = 72
    chunks.append(("" if first else " ") + remaining)
    return chunks


def event_lines(section: str, grouped: dict[str, list[Member]]) -> list[str]:
    start, end = SESSION_TIMES[section]
    slug = section_slug(section)
    roster = " | ".join(member.name for member in grouped[section])
    description = (
        f"Roster: {roster}. Source: team_training_schedule.xlsx. "
        "Time zone, location, organizer email, and attendee emails are not provided."
    )
    return [
        "BEGIN:VEVENT",
        f"UID:{section_uid(section)}",
        "DTSTAMP:20200501T100000Z",
        f"DTSTART:{TRAINING_DATE.strftime('%Y%m%d')}T{start.strftime('%H%M%S')}",
        f"DTEND:{TRAINING_DATE.strftime('%Y%m%d')}T{end.strftime('%H%M%S')}",
        f"SUMMARY:Team training - {section}",
        "STATUS:TENTATIVE",
        f"DESCRIPTION:{escape_ics(description)}",
        f"X-EVOCOWORK-SECTION:{section}",
        f"X-EVOCOWORK-ROSTER:{escape_ics(roster)}",
        "X-EVOCOWORK-TIME-BASIS:FLOATING_LOCAL_TIME",
        "X-EVOCOWORK-EVIDENCE-GAPS:REAL_EMAIL_TIMEZONE_LOCATION",
        "END:VEVENT",
    ]


def build_ics_bytes(sections: list[str], grouped: dict[str, list[Member]]) -> bytes:
    logical_lines = [
        "BEGIN:VCALENDAR",
        "VERSION:2.0",
        "PRODID:-//EvoCowork//OfficeBench 3-6 Training Coordination//EN",
        "CALSCALE:GREGORIAN",
        "METHOD:PUBLISH",
    ]
    for section in sections:
        logical_lines.extend(event_lines(section, grouped))
    logical_lines.append("END:VCALENDAR")
    physical_lines = [chunk for line in logical_lines for chunk in fold_ics_line(line)]
    return ("\r\n".join(physical_lines) + "\r\n").encode("utf-8")


def build_master_calendar(members: list[Member], output_path: Path) -> None:
    grouped = members_by_section(members)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_bytes(build_ics_bytes(list(SESSION_TIMES), grouped))


def serialize_message(message: EmailMessage) -> bytes:
    with tempfile.TemporaryFile() as handle:
        BytesGenerator(handle, policy=SMTP.clone(max_line_length=78), mangle_from_=False).flatten(message)
        handle.seek(0)
        return handle.read()


def build_mbox(members: list[Member], output_path: Path) -> None:
    grouped = members_by_section(members)
    records: list[bytes] = []
    for member in members:
        start, end = SESSION_TIMES[member.section]
        slug = section_slug(member.section)
        uid = section_uid(member.section)
        reserved_address = f"{member.name.lower()}@example.invalid"
        attachment_name = f"{slug}-training.ics"
        message = EmailMessage(policy=SMTP)
        message["From"] = "Training Coordinator <training-coordinator@example.invalid>"
        message["To"] = f"{member.name} <{reserved_address}>"
        message["Subject"] = f"Team training reminder - {member.section} - 2020-05-08"
        message["Date"] = "Fri, 01 May 2020 10:00:00 -0000"
        message["Message-ID"] = f"<{member.name.lower()}-training-20200508@evocowork.invalid>"
        message["X-EvoCowork-Delivery-Status"] = "BLOCKED_MISSING_REAL_EMAIL"
        message["X-EvoCowork-Source-Row"] = str(member.source_row)
        message["X-EvoCowork-Session-UID"] = uid
        body = (
            f"Hello {member.name},\n\n"
            "This is a controlled reminder draft for team training.\n"
            f"Training date: 2020-05-08\n"
            f"Section: {member.section}\n"
            f"Time: {start.strftime('%H:%M')}-{end.strftime('%H:%M')} (floating local time; time zone not provided)\n"
            "Location: not provided\n"
            f"Session UID: {uid}\n\n"
            "Delivery status: BLOCKED_MISSING_REAL_EMAIL. The source workbook does not provide a real email address; "
            "the .invalid address is a non-deliverable placeholder. This message is a draft and has not been sent.\n"
        )
        message.set_content(body, subtype="plain", charset="utf-8")
        message.add_attachment(
            build_ics_bytes([member.section], grouped),
            maintype="text",
            subtype="calendar",
            filename=attachment_name,
            disposition="attachment",
            params={"method": "PUBLISH", "charset": "utf-8"},
        )
        message.set_boundary(f"===============EvoCoworkTraining{member.source_row:02d}==")
        serialized = serialize_message(message)
        unix_from = b"From training-coordinator@example.invalid Fri May 01 10:00:00 2020\n"
        records.append(unix_from + serialized + b"\n")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_bytes(b"".join(records))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    members = parse_source(args.input_dir)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for name in (XLSX_NAME, DOCX_NAME, ICS_NAME, MBOX_NAME):
        existing = args.output_dir / name
        if existing.is_file():
            existing.unlink()
    build_workbook(members, args.output_dir / XLSX_NAME)
    build_document(members, args.output_dir / DOCX_NAME)
    build_master_calendar(members, args.output_dir / ICS_NAME)
    build_mbox(members, args.output_dir / MBOX_NAME)


if __name__ == "__main__":
    main()
