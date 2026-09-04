#!/usr/bin/env python3
"""OfficeBench 3-6 团队培训协调包的 v2 原子连续评分。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import mailbox
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import date, datetime, time
from email import policy
from email.parser import BytesParser
from email.utils import parseaddr
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from openpyxl import load_workbook


TASK_ID = "cowork-evo/officebench-3-6-team-training-coordination-package"
RUBRIC_VERSION = "officebench-3-6-team-training-coordination-v0.2"
JUDGE_VERSION = "team-training-atomic-package-v2.0"
SOURCE = "team_training_schedule.xlsx"
SOURCE_SHA256 = "a5628f45d9f48eccd8424634fda4d3c9048200571cc03db1e6d34cf9447498f4"
XLSX = "Team_Training_Coordination_Audit.xlsx"
DOCX = "Team_Training_Coordination_Brief.docx"
ICS = "Team_Training_Master.ics"
MBOX = "Team_Training_Reminders.mbox"
REQUIRED = (XLSX, DOCX, ICS, MBOX)
WEIGHTS = {"R001": 6, "R002": 10, "R003": 9, "R004": 10, "R005": 7, "R006": 10,
           "R007": 13, "R008": 10, "R009": 9, "R010": 6, "R011": 4, "R012": 6}
PENALTIES = {"P001": -12, "P002": -4}
TRAINING_DATE = date(2020, 5, 8)
ROWS = (
    (2, "Alice", "Section A", "9:00 AM - 11:00 AM"),
    (3, "Bob", "Section B", "11:00 AM - 1:00 PM"),
    (4, "Carol", "Section C", "2:00 PM - 4:00 PM"),
    (5, "David", "Section A", "9:00 AM - 11:00 AM"),
    (6, "Eva", "Section B", "11:00 AM - 1:00 PM"),
    (7, "Frank", "Section C", "2:00 PM - 4:00 PM"),
    (8, "Grace", "Section A", "9:00 AM - 11:00 AM"),
    (9, "Hank", "Section B", "11:00 AM - 1:00 PM"),
    (10, "Ivy", "Section C", "2:00 PM - 4:00 PM"),
    (11, "Jack", "Section A", "9:00 AM - 11:00 AM"),
    (12, "Kelly", "Section B", "11:00 AM - 1:00 PM"),
    (13, "Leo", "Section C", "2:00 PM - 4:00 PM"),
)
SESSION = {
    "Section A": (time(9), time(11), ("Alice", "David", "Grace", "Jack")),
    "Section B": (time(11), time(13), ("Bob", "Eva", "Hank", "Kelly")),
    "Section C": (time(14), time(16), ("Carol", "Frank", "Ivy", "Leo")),
}


def _load_v1():
    path = Path(__file__).with_name("judge_v1.py")
    spec = importlib.util.spec_from_file_location("team_training_judge_v1", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v1 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


V1 = _load_v1()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _task_root() -> Path:
    return Path(__file__).resolve().parents[1]


def _input_dir() -> Path:
    return Path("/app/input_files") if Path("/app/input_files").is_dir() else _task_root() / "data" / "input_files"


def _norm(value: Any) -> str:
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold()


def _same(actual: Any, expected: Any) -> bool:
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        try:
            return abs(float(actual) - float(expected)) <= 0.011
        except (TypeError, ValueError):
            return False
    if isinstance(expected, date) and not isinstance(expected, datetime):
        if isinstance(actual, datetime):
            actual = actual.date()
        return actual == expected
    if isinstance(expected, time):
        if isinstance(actual, datetime):
            actual = actual.time()
        return actual == expected
    return _norm(actual) == _norm(expected)


def _ratio(items: Iterable[tuple[bool, str]]) -> tuple[float, list[str], int, int]:
    findings = list(items)
    earned = sum(bool(ok) for ok, _ in findings)
    return earned / max(1, len(findings)), [label for ok, label in findings if not ok], earned, len(findings)


def _result(identifier: str, items: Iterable[tuple[bool, str]]) -> dict[str, Any]:
    score, failures, earned, total = _ratio(items)
    return {
        "criterion_id": identifier,
        "score": round(score, 6),
        "evidence": "全部任务专属原子检查通过" if not failures else "未通过：" + "；".join(failures[:24]),
        "method": "deterministic",
        "raw": {"earned_units": earned, "total_units": total, "failure_count": len(failures)},
    }


def _formula(value: Any, *tokens: str) -> bool:
    if not isinstance(value, str) or not value.startswith("="):
        return False
    normalized = re.sub(r"[\s'$]", "", value).upper()
    return all(re.sub(r"[\s'$]", "", token).upper() in normalized for token in tokens)


def _core(artifact: Path) -> tuple[bool, list[str], list[str]]:
    errors: list[str] = []
    hashes: list[str] = []
    if not artifact.is_dir():
        return False, ["输出目录不存在"], hashes
    for name in REQUIRED:
        path = artifact / name
        if not path.is_file() or path.stat().st_size == 0:
            errors.append(f"核心交付缺失或为空：{name}")
            continue
        try:
            hashes.append(_sha256(path))
            if name == XLSX:
                if not zipfile.is_zipfile(path):
                    raise ValueError("不是有效 OOXML ZIP")
                book = load_workbook(path, read_only=True, data_only=False)
                if not book.sheetnames:
                    raise ValueError("工作簿无工作表")
                book.close()
            elif name == DOCX:
                if not zipfile.is_zipfile(path):
                    raise ValueError("不是有效 OOXML ZIP")
                document = Document(path)
                if not document.paragraphs and not document.tables:
                    raise ValueError("Word 无内容")
            elif name == ICS:
                parsed = V1.parse_calendar_bytes(path.read_bytes())
                if not parsed.get("events"):
                    raise ValueError("ICS 无事件")
            else:
                if not path.read_bytes().startswith(b"From "):
                    raise ValueError("不是 Unix mbox")
                messages = list(mailbox.mbox(path, factory=lambda h: BytesParser(policy=policy.default).parse(h), create=False))
                if not messages:
                    raise ValueError("MBOX 无邮件")
        except Exception as exc:
            errors.append(f"核心交付无法解析：{name}: {type(exc).__name__}: {exc}")
    source = _input_dir() / SOURCE
    if not source.is_file() or _sha256(source) != SOURCE_SHA256:
        errors.append("冻结输入缺失或 SHA256 错误")
    return not errors, errors, hashes


def _fatal(errors: list[str], hashes: list[str]) -> dict[str, Any]:
    criterion_results = [_result(identifier, [(False, error) for error in errors]) for identifier in WEIGHTS]
    criterion_results.extend(_result(identifier, [(True, "未触发")]) for identifier in PENALTIES)
    return {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes or ["0" * 64], "criterion_results": criterion_results,
        "reward": 0.0, "passed": False, "hurdle_failures": ["R001"], "penalties": [],
    }


def _recalculate(path: Path) -> tuple[dict[str, list[list[Any]]] | None, str]:
    with tempfile.TemporaryDirectory(prefix="training-v2-recalc-") as raw:
        root = Path(raw); source = root / "source"; output = root / "output"; profile = root / "profile"
        source.mkdir(); output.mkdir(); copied = source / path.name; shutil.copy2(path, copied)
        env = os.environ.copy(); env["HOME"] = str(root / "home"); env["NO_PROXY"] = "localhost,127.0.0.1,::1"
        try:
            completed = subprocess.run(
                ["libreoffice", "--headless", f"-env:UserInstallation=file://{profile}", "--convert-to", "xlsx", "--outdir", str(output), str(copied)],
                stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300, env=env, check=False,
            )
        except Exception as exc:
            return None, str(exc)
        converted = output / path.name
        if completed.returncode != 0 or not converted.is_file():
            return None, f"退出={completed.returncode}；日志={completed.stdout[-400:]}"
        try:
            book = load_workbook(converted, data_only=True)
            values = {sheet.title: [[cell.value for cell in row] for row in sheet.iter_rows()] for sheet in book.worksheets}
            book.close()
            return values, completed.stdout[-400:]
        except Exception as exc:
            return None, str(exc)


def _value(values: dict[str, list[list[Any]]] | None, sheet: str, row: int, column: int) -> Any:
    try:
        return values[sheet][row - 1][column - 1] if values is not None else None
    except (KeyError, IndexError):
        return None


def _ooxml_danger(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as archive:
        return [name for name in archive.namelist() if "vbaproject" in name.casefold() or name.startswith("xl/externalLinks/") or "/embeddings/" in name or "oleobject" in name.casefold()]


def _workbook(artifact: Path) -> tuple[dict[str, list[tuple[bool, str]]], dict[str, Any], list[str]]:
    findings = {key: [] for key in ("R002", "R003", "R004", "R005", "R011", "R012")}
    context: dict[str, Any] = {"names": [], "sections": {}, "blocked": 0}
    dangers: list[str] = []
    path = artifact / XLSX
    try:
        dangers = _ooxml_danger(path)
        book = load_workbook(path, data_only=False)
        values, recalc_log = _recalculate(path)
    except Exception as exc:
        for key in findings:
            findings[key].append((False, f"Excel 检查失败：{exc}"))
        return findings, context, dangers
    expected_sheets = ["Roster_Control", "Session_Control", "Reminder_Control", "Delivery_Summary"]
    for index, name in enumerate(expected_sheets):
        findings["R012"].append((index < len(book.sheetnames) and book.sheetnames[index] == name, f"工作表第{index + 1}项错误"))

    roster_headers = ["Source Row", "Name", "Section", "Source Time", "Start", "End", "Duration (Hours)", "Training Date", "Reserved Address", "Delivery Status", "Session UID"]
    roster = book["Roster_Control"] if "Roster_Control" in book.sheetnames else None
    for column, expected in enumerate(roster_headers, 1):
        findings["R002"].append((roster is not None and _same(roster.cell(1, column).value, expected), f"Roster 表头第{column}列错误"))
    for row, (source_row, name, section, source_time) in enumerate(ROWS, 2):
        start, end, _ = SESSION[section]
        expected = (source_row, name, section, source_time, start, end, TRAINING_DATE, "BLOCKED_MISSING_REAL_EMAIL")
        cells = (1, 2, 3, 4, 5, 6, 8, 10)
        for column, target in zip(cells, expected, strict=True):
            actual = roster.cell(row, column).value if roster else None
            findings["R002"].append((_same(actual, target), f"Roster 第{row}行第{column}列错误"))
        context["names"].append(name if roster and _same(roster.cell(row, 2).value, name) else roster.cell(row, 2).value if roster else None)
        context["blocked"] += int(bool(roster and roster.cell(row, 10).value == "BLOCKED_MISSING_REAL_EMAIL"))
        formula_targets = [
            (7, ("ROUND", "MOD", f"F{row}-E{row}"), 2),
            (9, ("LOWER", f"B{row}", "EXAMPLE.INVALID"), f"{name.lower()}@example.invalid"),
            (11, ("LOWER", "SUBSTITUTE", f"C{row}"), f"{section.lower().replace(' ', '-')}-training-20200508@evocowork.invalid"),
        ]
        for column, tokens, expected_value in formula_targets:
            formula = roster.cell(row, column).value if roster else None
            findings["R004"].append((_formula(formula, *tokens), f"Roster!{chr(64 + column)}{row} 公式错误"))
            findings["R004"].append((_same(_value(values, "Roster_Control", row, column), expected_value), f"Roster!{chr(64 + column)}{row} 重算错误"))

    session_headers = ["Section", "Training Date", "Start", "End", "Duration (Hours)", "Member Count", "Members", "Previous End", "Gap (Minutes)", "Overlap Flag", "Session UID", "Calendar Status"]
    sessions = book["Session_Control"] if "Session_Control" in book.sheetnames else None
    for column, expected in enumerate(session_headers, 1):
        findings["R003"].append((sessions is not None and _same(sessions.cell(1, column).value, expected), f"Session 表头第{column}列错误"))
    for row, (section, (start, end, names)) in enumerate(SESSION.items(), 2):
        expected = (section, TRAINING_DATE, start, end, " | ".join(names), f"{section.lower().replace(' ', '-')}-training-20200508@evocowork.invalid", "TENTATIVE")
        for column, target in zip((1, 2, 3, 4, 7, 11, 12), expected, strict=True):
            actual = sessions.cell(row, column).value if sessions else None
            findings["R003"].append((_same(actual, target), f"Session {section} 第{column}列错误"))
        context["sections"][section] = {
            "start": sessions.cell(row, 3).value if sessions else None,
            "end": sessions.cell(row, 4).value if sessions else None,
            "members": sessions.cell(row, 7).value if sessions else None,
            "uid": sessions.cell(row, 11).value if sessions else None,
        }
        formula_specs = [(5, ("ROUND", "MOD"), 2), (6, ("COUNTIF", "ROSTER_CONTROL", f"A{row}"), 4)]
        if row > 2:
            formula_specs.extend([(9, ("ROUND", "1440"), 0 if row == 3 else 60), (10, ("IF", f"C{row}<H{row}"), "NO")])
        for column, tokens, expected_value in formula_specs:
            formula = sessions.cell(row, column).value if sessions else None
            findings["R004"].append((_formula(formula, *tokens), f"Session!{chr(64 + column)}{row} 公式错误"))
            findings["R004"].append((_same(_value(values, "Session_Control", row, column), expected_value), f"Session!{chr(64 + column)}{row} 重算错误：{recalc_log}"))
    findings["R004"].append((_same(_value(values, "Session_Control", 2, 9), None), "Section A 首行 Gap 应为空"))
    findings["R004"].append((_same(sessions.cell(2, 10).value if sessions else None, "N/A"), "Section A Overlap 应为 N/A"))

    reminder_headers = ["Source Row", "Name", "Section", "Reserved Recipient", "Subject", "Message ID", "Calendar Attachment", "Session UID", "Delivery Status"]
    reminders = book["Reminder_Control"] if "Reminder_Control" in book.sheetnames else None
    for column, expected in enumerate(reminder_headers, 1):
        findings["R005"].append((reminders is not None and _same(reminders.cell(1, column).value, expected), f"Reminder 表头第{column}列错误"))
    for row, (source_row, name, section, _) in enumerate(ROWS, 2):
        slug = section.lower().replace(" ", "-")
        expected = (source_row, name, section, f"{name.lower()}@example.invalid", f"Team training reminder - {section} - 2020-05-08", f"<{name.lower()}-training-20200508@evocowork.invalid>", f"{slug}-training.ics", f"{slug}-training-20200508@evocowork.invalid", "BLOCKED_MISSING_REAL_EMAIL")
        for column, target in enumerate(expected, 1):
            actual = _value(values, "Reminder_Control", row, column) if column in (4, 5) else (reminders.cell(row, column).value if reminders else None)
            findings["R005"].append((_same(actual, target), f"Reminder {name} 第{column}列错误"))
        for column, tokens in ((4, ("LOWER", f"B{row}", "EXAMPLE.INVALID")), (5, ("TEAMTRAININGREMINDER", f"C{row}", "2020-05-08"))):
            findings["R004"].append((_formula(reminders.cell(row, column).value if reminders else None, *tokens), f"Reminder!{chr(64 + column)}{row} 公式错误"))

    summary = book["Delivery_Summary"] if "Delivery_Summary" in book.sheetnames else None
    summary_targets = (("Source Rows", 12), ("Sections", 3), ("Reminder Messages", 12), ("Blocked Drafts", 12), ("Calendar Events", 3), ("Scheduled Hours", 6), ("Inter-session Gap (Minutes)", 60), ("Overlap Count", 0))
    for row, (metric, expected) in enumerate(summary_targets, 2):
        findings["R004"].append((_same(summary.cell(row, 1).value if summary else None, metric), f"Summary 第{row}行指标名错误"))
        formula = summary.cell(row, 2).value if summary else None
        findings["R004"].append((isinstance(formula, str) and formula.startswith("="), f"Summary!B{row} 不是公式"))
        findings["R004"].append((_same(_value(values, "Delivery_Summary", row, 2), expected), f"Summary!B{row} 重算错误"))
        findings["R004"].append((_same(_value(values, "Delivery_Summary", row, 4), "PASS"), f"Summary!D{row} 未重算为 PASS"))

    tables = {"Roster_Control": "TrainingRosterControlTable", "Session_Control": "TrainingSessionControlTable", "Reminder_Control": "TrainingReminderControlTable", "Delivery_Summary": "TrainingDeliverySummaryTable"}
    for sheet_name, table_name in tables.items():
        sheet = book[sheet_name] if sheet_name in book.sheetnames else None
        findings["R012"].extend([
            (sheet is not None and table_name in sheet.tables, f"{sheet_name} 原生表缺失"),
            (sheet is not None and sheet.freeze_panes == "A2", f"{sheet_name} 未冻结首行"),
            (sheet is not None and sheet.auto_filter.ref is not None, f"{sheet_name} 缺筛选"),
        ])
    findings["R012"].append((summary is not None and len(summary._charts) == 1, "Delivery_Summary 缺可编辑柱状图"))
    findings["R012"].append((sum(len(sheet.conditional_formatting) for sheet in book.worksheets) >= 5, "条件格式不足"))
    findings["R011"].append((not dangers, f"Excel 含危险对象：{dangers}"))
    book.close()
    return findings, context, dangers


def _event_findings(event: dict[str, str] | None, section: str) -> list[tuple[bool, str]]:
    start, end, names = SESSION[section]
    roster = " | ".join(names); slug = section.lower().replace(" ", "-")
    expected = {
        "UID": f"{slug}-training-20200508@evocowork.invalid", "DTSTAMP": "20200501T100000Z",
        "DTSTART": f"20200508T{start.strftime('%H%M%S')}", "DTEND": f"20200508T{end.strftime('%H%M%S')}",
        "SUMMARY": f"Team training - {section}", "STATUS": "TENTATIVE", "X-EVOCOWORK-SECTION": section,
        "X-EVOCOWORK-ROSTER": roster, "X-EVOCOWORK-TIME-BASIS": "FLOATING_LOCAL_TIME",
        "X-EVOCOWORK-EVIDENCE-GAPS": "REAL_EMAIL_TIMEZONE_LOCATION",
    }
    findings: list[tuple[bool, str]] = []
    for key, target in expected.items():
        findings.append((event is not None and _same(event.get(key), target), f"{section} 事件 {key} 错误"))
    description = event.get("DESCRIPTION", "") if event else ""
    for required in (roster, SOURCE, "Time zone", "location", "not provided"):
        findings.append((required.casefold() in description.casefold(), f"{section} DESCRIPTION 缺 {required}"))
    for banned in ("LOCATION", "ORGANIZER", "ATTENDEE", "GEO"):
        findings.append((event is None or banned not in event, f"{section} 伪造 {banned}"))
    findings.append((event is None or not any("TZID=" in str(value).upper() for value in event.values()), f"{section} 伪造 TZID"))
    return findings


def _master(artifact: Path) -> tuple[list[tuple[bool, str]], dict[str, dict[str, str]], list[tuple[bool, str]], int]:
    findings: list[tuple[bool, str]] = []; boundary: list[tuple[bool, str]] = []; event_map: dict[str, dict[str, str]] = {}; invented = 0
    try:
        parsed = V1.parse_calendar_bytes((artifact / ICS).read_bytes()); events = parsed["events"]
        findings.append((len(events) == 3, f"Master ICS 事件数错误：{len(events)}"))
        for event in events:
            section = event.get("X-EVOCOWORK-SECTION", "")
            if section in SESSION and section not in event_map:
                event_map[section] = event
        findings.append((set(event_map) == set(SESSION), "Master ICS 未恰好覆盖 A/B/C"))
        for section in SESSION:
            atom = _event_findings(event_map.get(section), section); findings.extend(atom[:15]); boundary.extend(atom[15:])
        extra_components = any(line.startswith(("BEGIN:VALARM", "BEGIN:VTIMEZONE")) for line in parsed["lines"])
        boundary.append((not extra_components, "Master ICS 增加无来源时区或提醒"))
    except Exception as exc:
        findings.append((False, f"Master ICS 语义解析失败：{exc}"))
    invented = sum(not ok for ok, _ in boundary)
    return findings, event_map, boundary, invented


def _mbox(artifact: Path, master: dict[str, dict[str, str]]) -> tuple[list[tuple[bool, str]], list[tuple[bool, str]], dict[str, Any], list[tuple[bool, str]], int]:
    envelope: list[tuple[bool, str]] = []; attachments: list[tuple[bool, str]] = []; boundary: list[tuple[bool, str]] = []
    context: dict[str, Any] = {"names": [], "uids": [], "recipients": []}; invented = 0
    messages = list(mailbox.mbox(artifact / MBOX, factory=lambda h: BytesParser(policy=policy.default).parse(h), create=False))
    envelope.append((len(messages) == 12, f"MBOX 邮件数错误：{len(messages)}"))
    by_name: dict[str, Any] = {}
    for message in messages:
        name, address = parseaddr(str(message.get("To", "")))
        envelope.append((bool(name) and name not in by_name, f"收件人姓名缺失或重复：{name!r}"))
        if name and name not in by_name:
            by_name[name] = message
        boundary.extend([(address.endswith("@example.invalid"), f"{name} 收件地址不是 .invalid"), (not message.get("Cc") and not message.get("Bcc"), f"{name} 增加 Cc/Bcc")])
    for source_row, name, section, _ in ROWS:
        message = by_name.get(name); start, end, _ = SESSION[section]; slug = section.lower().replace(" ", "-")
        uid = f"{slug}-training-20200508@evocowork.invalid"; address = f"{name.lower()}@example.invalid"
        actual_address = parseaddr(str(message.get("To", "")))[1] if message else ""
        targets = [
            (message is not None, f"缺少 {name} 的提醒"),
            (message is not None and message.is_multipart(), f"{name} 不是 multipart"),
            (_same(actual_address, address), f"{name} 收件地址错误"),
            (_same(message.get("Subject") if message else None, f"Team training reminder - {section} - 2020-05-08"), f"{name} 主题错误"),
            (_same(message.get("Message-ID") if message else None, f"<{name.lower()}-training-20200508@evocowork.invalid>"), f"{name} Message-ID 错误"),
            (_same(message.get("X-EvoCowork-Delivery-Status") if message else None, "BLOCKED_MISSING_REAL_EMAIL"), f"{name} 阻塞头错误"),
            (_same(message.get("X-EvoCowork-Source-Row") if message else None, str(source_row)), f"{name} 源行头错误"),
            (_same(message.get("X-EvoCowork-Session-UID") if message else None, uid), f"{name} Session UID 头错误"),
        ]
        body_part = message.get_body(preferencelist=("plain",)) if message else None
        body = body_part.get_content() if body_part else ""
        for required in (f"Hello {name}", "2020-05-08", section, f"{start.strftime('%H:%M')}-{end.strftime('%H:%M')}", "floating local time", "time zone not provided", "Location: not provided", uid, "BLOCKED_MISSING_REAL_EMAIL", ".invalid", "draft", "has not been sent"):
            targets.append((required.casefold() in body.casefold(), f"{name} 正文缺 {required}"))
        envelope.extend(targets)
        parts = list(message.iter_attachments()) if message else []
        attachments.append((len(parts) == 1, f"{name} 日历附件数错误：{len(parts)}"))
        if len(parts) == 1:
            part = parts[0]
            attachments.extend([
                (_same(part.get_content_type(), "text/calendar"), f"{name} 附件 MIME 错误"),
                (_same(part.get_filename(), f"{slug}-training.ics"), f"{name} 附件名错误"),
            ])
            try:
                parsed = V1.parse_calendar_bytes(part.get_payload(decode=True) or b""); events = parsed["events"]
                attachments.append((len(events) == 1, f"{name} 附件事件数错误"))
                event = events[0] if len(events) == 1 else None
                atom = _event_findings(event, section); attachments.extend(atom[:15]); boundary.extend(atom[15:])
                compare = ("UID", "DTSTART", "DTEND", "SUMMARY", "STATUS", "DESCRIPTION", "X-EVOCOWORK-SECTION", "X-EVOCOWORK-ROSTER", "X-EVOCOWORK-TIME-BASIS", "X-EVOCOWORK-EVIDENCE-GAPS")
                attachments.append((event is not None and master.get(section) is not None and all(_same(event.get(key), master[section].get(key)) for key in compare), f"{name} 附件与 Master ICS 不一致"))
                context["uids"].append(event.get("UID") if event else None)
            except Exception as exc:
                attachments.append((False, f"{name} 附件无法解析：{exc}"))
        context["names"].append(name if message else None); context["recipients"].append(actual_address)
    expected_names = [row[1] for row in ROWS]
    envelope.append((set(by_name) == set(expected_names), "MBOX 收件人集合与源表不一致"))
    invented = sum(not ok for ok, _ in boundary)
    return envelope, attachments, context, boundary, invented


def _document(artifact: Path) -> tuple[list[tuple[bool, str]], dict[str, Any], list[tuple[bool, str]], list[str]]:
    findings: list[tuple[bool, str]] = []; boundary: list[tuple[bool, str]] = []; dangers: list[str] = []
    path = artifact / DOCX; document = Document(path); dangers = _ooxml_danger(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    text = "\n".join(parts)
    required = ["Team Training Coordination Brief", "2020-05-08", "12 members", "3 sections", "12 reminder drafts blocked", "Section A", "09:00-11:00", "Section B", "11:00-13:00", "Section C", "14:00-16:00", "60-minute inter-session gap", "exactly two hours", "BLOCKED_MISSING_REAL_EMAIL", ".invalid", "floating local time", "No TZID", "No LOCATION", "No ORGANIZER or ATTENDEE", "not sent", "remain tentative"]
    for item in required:
        findings.append((item.casefold() in text.casefold(), f"Word 缺少 {item}"))
    for _, name, section, source_time in ROWS:
        for item in (name, section, source_time, f"{name.lower()}@example.invalid"):
            findings.append((item.casefold() in text.casefold(), f"Word 缺 {name} 的 {item}"))
    findings.extend([(len(document.tables) == 4, f"Word 原生表格数错误：{len(document.tables)}"), (not document.inline_shapes, "Word 使用图片替代表格")])
    expected_tables = [
        [
            ["Section", "Date", "Floating local time", "Members", "Calendar status"],
            *[[section, "2020-05-08", f"{start.strftime('%H:%M')}-{end.strftime('%H:%M')}", " | ".join(names), "TENTATIVE"] for section, (start, end, names) in SESSION.items()],
        ],
        [
            ["Missing source field", "Required handling", "Control outcome"],
            ["Real recipient emails", "Use non-deliverable .invalid placeholders only", "12 drafts BLOCKED_MISSING_REAL_EMAIL"],
            ["Time zone", "Keep DTSTART/DTEND as floating local time", "No TZID and no UTC conversion"],
            ["Location", "State that location is not provided", "No LOCATION property"],
            ["Organizer/attendee mailboxes", "Do not fabricate identities or addresses", "No ORGANIZER or ATTENDEE property"],
        ],
        [
            ["Row", "Name", "Section", "Source time", "Reserved recipient", "Status"],
            *[[str(source_row), name, section, source_time, f"{name.lower()}@example.invalid", "BLOCKED"] for source_row, name, section, source_time in ROWS],
        ],
        [
            ["Artifact", "Required content", "Status"],
            [XLSX, "4 native audit tables, formulas and editable section chart", "READY"],
            [DOCX, "Schedule, roster, evidence gaps and controls", "READY"],
            [ICS, "3 tentative floating-time events with stable UIDs", "READY"],
            [MBOX, "12 multipart drafts with section calendar attachments", "BLOCKED"],
        ],
    ]
    for table_index, expected_rows in enumerate(expected_tables):
        table = document.tables[table_index] if table_index < len(document.tables) else None
        findings.append((table is not None and len(table.rows) == len(expected_rows), f"Word 表{table_index + 1}行数错误"))
        for row_index, expected_row in enumerate(expected_rows):
            for column_index, expected in enumerate(expected_row):
                actual = table.rows[row_index].cells[column_index].text if table is not None and row_index < len(table.rows) and column_index < len(table.rows[row_index].cells) else None
                findings.append((_same(actual, expected), f"Word 表{table_index + 1} 第{row_index + 1}行第{column_index + 1}列错误"))
    section = document.sections[0]
    findings.append((abs(section.page_width.cm - 21.0) <= 0.2 and abs(section.page_height.cm - 29.7) <= 0.2, "Word 页面不是 A4 纵向"))
    safe_text = re.sub(r"[A-Za-z0-9._%+-]+@(?:example|evocowork)\.invalid", "", text, flags=re.IGNORECASE)
    for pattern, label in ((r"@[A-Za-z0-9.-]+\.(?:com|edu|org|net)\b", "真实邮箱"), (r"\b(?:GMT|PST|EST|CST)\b", "时区"), (r"\bRoom\s+\w+\b", "地点"), (r"\b(?:has|have|was|were) been sent\b", "已发送声明")):
        boundary.append((not re.search(pattern, safe_text, flags=re.IGNORECASE), f"Word 编造{label}"))
    boundary.append((not dangers, f"Word 含危险对象：{dangers}"))
    return findings, {"text": text}, boundary, dangers


def _render(path: Path) -> tuple[bool, int, str]:
    with tempfile.TemporaryDirectory(prefix="training-v2-render-") as raw:
        root = Path(raw); output = root / "output"; profile = root / "profile"; output.mkdir()
        env = os.environ.copy(); env["HOME"] = str(root / "home"); env["NO_PROXY"] = "localhost,127.0.0.1,::1"
        completed = subprocess.run(["libreoffice", "--headless", f"-env:UserInstallation=file://{profile}", "--convert-to", "pdf", "--outdir", str(output), str(path)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300, env=env, check=False)
        pdf = output / f"{path.stem}.pdf"
        if completed.returncode != 0 or not pdf.is_file() or pdf.stat().st_size < 2000:
            return False, 0, completed.stdout[-400:]
        info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=60, check=False)
        match = re.search(r"^Pages:\s+(\d+)\s*$", info.stdout, flags=re.MULTILINE)
        return bool(match), int(match.group(1)) if match else 0, info.stdout[-400:]


def evaluate(artifact: Path, scoring_version: str | None = None) -> dict[str, Any]:
    if (scoring_version or os.environ.get("SCORING_VERSION", "v2")).casefold() == "v1":
        return V1.evaluate(artifact)
    core_ok, core_errors, hashes = _core(artifact)
    if not core_ok:
        return _fatal(core_errors, hashes)
    criteria: dict[str, list[tuple[bool, str]]] = {key: [] for key in WEIGHTS}
    criteria["R001"].append((True, "四份核心交付与冻结输入均可解析"))
    workbook_findings, workbook_context, excel_dangers = _workbook(artifact)
    for key, items in workbook_findings.items():
        criteria[key].extend(items)
    master_findings, master_context, master_boundary, master_invented = _master(artifact)
    criteria["R006"].extend(master_findings); criteria["R011"].extend(master_boundary)
    mail_findings, attachment_findings, mail_context, mail_boundary, mail_invented = _mbox(artifact, master_context)
    criteria["R007"].extend(mail_findings); criteria["R008"].extend(attachment_findings); criteria["R011"].extend(mail_boundary)
    document_findings, document_context, document_boundary, doc_dangers = _document(artifact)
    criteria["R009"].extend(document_findings); criteria["R011"].extend(document_boundary)

    expected_names = [row[1] for row in ROWS]
    expected_uids = [f"{row[2].lower().replace(' ', '-')}-training-20200508@evocowork.invalid" for row in ROWS]
    criteria["R010"].extend([
        (workbook_context.get("names") == expected_names, "Excel 名单顺序与源表不一致"),
        (set(master_context) == set(SESSION), "Master ICS Section 集合不一致"),
        (mail_context.get("names") == expected_names, "MBOX 名单顺序与源表不一致"),
        (mail_context.get("uids") == expected_uids, "MBOX 日历 UID 顺序不一致"),
        (workbook_context.get("blocked") == 12, "Excel 阻塞人数不是12"),
        (all(name in document_context.get("text", "") for name in expected_names), "Word 与名单不一致"),
    ])
    for section, (start, end, names) in SESSION.items():
        excel_section = workbook_context.get("sections", {}).get(section, {})
        master = master_context.get(section, {})
        criteria["R010"].extend([
            (_same(excel_section.get("uid"), master.get("UID")), f"{section} Excel/ICS UID 不一致"),
            (_same(master.get("DTSTART"), f"20200508T{start.strftime('%H%M%S')}"), f"{section} 跨文件开始时间错误"),
            (_same(master.get("DTEND"), f"20200508T{end.strftime('%H%M%S')}"), f"{section} 跨文件结束时间错误"),
            (_same(excel_section.get("members"), " | ".join(names)), f"{section} Excel 名单错误"),
        ])

    for name, expected_pages in ((XLSX, 4), (DOCX, 2)):
        try:
            ok, pages, log = _render(artifact / name)
            criteria["R012"].append((ok, f"{name} 无法渲染：{log}"))
            criteria["R012"].append((pages == expected_pages, f"{name} 页数错误：{pages}"))
        except Exception as exc:
            criteria["R012"].append((False, f"{name} 渲染异常：{exc}"))

    actual_files = {path.name for path in artifact.iterdir() if path.is_file()}
    extra = sorted(actual_files - set(REQUIRED))
    extra_entries = sorted(path.name for path in artifact.iterdir() if not path.is_file())
    p1_bad = master_invented + mail_invented + sum(not ok for ok, _ in document_boundary)
    p1_score = min(1.0, p1_bad / 4.0)
    p2_bad = bool(extra or extra_entries or excel_dangers or doc_dangers)
    positive_results = [_result(identifier, criteria[identifier]) for identifier in WEIGHTS]
    base = sum(WEIGHTS[item["criterion_id"]] * float(item["score"]) for item in positive_results) / 100.0
    penalty_amounts = []
    if p1_score:
        penalty_amounts.append({"id": "P001", "amount": round(0.12 * p1_score, 6), "reason": "无来源邮箱/时区/地点/发送声明"})
    if p2_bad:
        penalty_amounts.append({"id": "P002", "amount": 0.04, "reason": f"额外交付或危险对象：extra={extra + extra_entries}"})
    total_penalty = sum(item["amount"] for item in penalty_amounts)
    reward = round(max(0.0, base - total_penalty), 6)
    positive_results.extend([
        {"criterion_id": "P001", "score": round(p1_score, 6), "evidence": "未触发" if not p1_score else "检测到无来源声明", "method": "deterministic", "raw": {"bad_units": p1_bad}},
        {"criterion_id": "P002", "score": float(p2_bad), "evidence": "未触发" if not p2_bad else "检测到额外交付或危险对象", "method": "deterministic", "raw": {"extra": extra + extra_entries}},
    ])
    return {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes, "criterion_results": positive_results, "base_score": round(base, 6),
        "reward": reward, "passed": reward >= 0.8, "hurdle_failures": [], "penalties": penalty_amounts,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    parser.add_argument("--scoring-version", choices=("v1", "v2"), default=os.environ.get("SCORING_VERSION", "v2"))
    args = parser.parse_args()
    try:
        result = evaluate(args.artifact_dir, args.scoring_version)
    except Exception as exc:
        result = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
                  "reward": 0.0, "passed": False, "fatal_error": f"{type(exc).__name__}: {exc}", "hurdle_failures": ["R001"]}
    args.result.parent.mkdir(parents=True, exist_ok=True); args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{float(result['reward']):.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
