#!/usr/bin/env python3
"""团队培训协调包的 v2 任务专属分层负例。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import mailbox
import shutil
import sys
import tempfile
from email import policy
from email.generator import BytesGenerator
from email.message import EmailMessage
from email.parser import BytesParser
from pathlib import Path
from typing import Callable

from docx import Document
from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
INPUTS = ROOT / "environment" / "input_files"
REQUIRED = ("Team_Training_Coordination_Audit.xlsx", "Team_Training_Coordination_Brief.docx", "Team_Training_Master.ics", "Team_Training_Reminders.mbox")
Mutator = Callable[[Path], None]


def _load_judge():
    path = Path(__file__).with_name("judge.py")
    spec = importlib.util.spec_from_file_location("training_v2_judge_for_negatives", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v2 Judge")
    module = importlib.util.module_from_spec(spec); sys.modules[spec.name] = module; spec.loader.exec_module(module)
    return module


JUDGE = _load_judge()


def _copy(target: Path) -> None:
    target.mkdir(parents=True, exist_ok=True)
    for name in REQUIRED:
        shutil.copy2(GOLDEN / name, target / name)


def _excel(sheet: str, cell: str, value: object) -> Mutator:
    def apply(root: Path) -> None:
        path = root / REQUIRED[0]; book = load_workbook(path); book[sheet][cell] = value; book.save(path); book.close()
    return apply


def _remove_chart(root: Path) -> None:
    path = root / REQUIRED[0]; book = load_workbook(path); book["Delivery_Summary"]._charts = []; book.save(path); book.close()


def _remove_freeze(root: Path) -> None:
    path = root / REQUIRED[0]; book = load_workbook(path); book["Roster_Control"].freeze_panes = None; book.save(path); book.close()


def _word_replace(old: str, new: str) -> Mutator:
    def apply(root: Path) -> None:
        path = root / REQUIRED[1]; document = Document(path)
        locations = list(document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    locations.extend(cell.paragraphs)
        for paragraph in locations:
            if old in paragraph.text:
                paragraph.text = paragraph.text.replace(old, new, 1); document.save(path); return
        raise AssertionError(f"Word 未找到：{old}")
    return apply


def _word_remove_table(root: Path) -> None:
    path = root / REQUIRED[1]; document = Document(path); table = document.tables[-1]; table._element.getparent().remove(table._element); document.save(path)


def _ics_replace(old: bytes, new: bytes) -> Mutator:
    def apply(root: Path) -> None:
        path = root / REQUIRED[2]; raw = path.read_bytes(); assert old in raw; path.write_bytes(raw.replace(old, new, 1))
    return apply


def _drop_master_event(root: Path) -> None:
    path = root / REQUIRED[2]; raw = path.read_bytes()
    start = raw.find(b"BEGIN:VEVENT\r\n", raw.find(b"BEGIN:VEVENT\r\n") + 1)
    end = raw.find(b"END:VEVENT\r\n", start) + len(b"END:VEVENT\r\n")
    assert start >= 0 and end > start
    path.write_bytes(raw[:start] + raw[end:])


def _read_messages(path: Path):
    return list(mailbox.mbox(path, factory=lambda h: BytesParser(policy=policy.default).parse(h), create=False))


def _write_messages(path: Path, messages: list[EmailMessage]) -> None:
    with tempfile.TemporaryDirectory(prefix="training-v2-mbox-") as raw:
        target = Path(raw) / "mail.mbox"; box = mailbox.mbox(target, create=True); box.lock()
        try:
            for message in messages:
                if message.get_unixfrom() is None:
                    message.set_unixfrom("From training-coordinator@example.invalid Fri May 01 10:00:00 2020")
                box.add(message)
            box.flush()
        finally:
            box.unlock(); box.close()
        shutil.copy2(target, path)


def _mail(mutator: Callable[[list[EmailMessage]], None]) -> Mutator:
    def apply(root: Path) -> None:
        path = root / REQUIRED[3]; messages = _read_messages(path); mutator(messages); _write_messages(path, messages)
    return apply


def _replace_header(index: int, header: str, value: str) -> Mutator:
    def mutate(messages: list[EmailMessage]) -> None:
        messages[index].replace_header(header, value)
    return _mail(mutate)


def _add_cc(messages: list[EmailMessage]) -> None:
    messages[0]["Cc"] = "Manager <manager@example.invalid>"


def _body_replace(index: int, old: str, new: str) -> Mutator:
    def mutate(messages: list[EmailMessage]) -> None:
        part = messages[index].get_body(preferencelist=("plain",)); text = part.get_content(); assert old in text
        part.set_content(text.replace(old, new, 1), subtype="plain", charset="utf-8")
    return _mail(mutate)


def _drop_attachment(messages: list[EmailMessage]) -> None:
    payload = messages[0].get_payload(); assert isinstance(payload, list) and len(payload) >= 2; messages[0].set_payload(payload[:-1])


def _attachment_time(messages: list[EmailMessage]) -> None:
    part = next(messages[0].iter_attachments()); raw = (part.get_payload(decode=True) or b"").replace(b"DTSTART:20200508T090000", b"DTSTART:20200508T100000", 1)
    filename = part.get_filename(); part.set_content(raw, maintype="text", subtype="calendar", cte="base64", params={"method": "PUBLISH", "charset": "utf-8"}); part.add_header("Content-Disposition", "attachment", filename=filename)


def _drop_message(messages: list[EmailMessage]) -> None:
    messages.pop()


def _duplicate_recipient(messages: list[EmailMessage]) -> None:
    messages[1].replace_header("To", "Alice <alice@example.invalid>")


def _moderate_communication_domain(root: Path) -> None:
    # 仍保持四份文件可打开，但把整个日历/邮件域改成错误时段、真实地址和已发送状态。
    ics = root / REQUIRED[2]; raw = ics.read_bytes().replace(b"20200508T090000", b"20200508T100000").replace(b"20200508T110000", b"20200508T120000").replace(b"20200508T130000", b"20200508T140000").replace(b"20200508T140000", b"20200508T150000").replace(b"20200508T160000", b"20200508T170000")
    raw = raw.replace(b"STATUS:TENTATIVE", b"STATUS:CONFIRMED").replace(b"END:VEVENT", b"LOCATION:Room 7\r\nEND:VEVENT")
    ics.write_bytes(raw)
    messages = _read_messages(root / REQUIRED[3])
    for message in messages:
        name = str(message.get("To", "")).split(" <", 1)[0]
        message.replace_header("To", f"{name} <{name.lower()}@company.com>")
        message.replace_header("X-EvoCowork-Delivery-Status", "SENT")
        body = message.get_body(preferencelist=("plain",)); body.set_content(f"Hello {name},\n\nYour confirmed training was sent. Location: Room 7.\n", subtype="plain", charset="utf-8")
    _write_messages(root / REQUIRED[3], messages)


def _severe_valid_incomplete(root: Path) -> None:
    book = Workbook(); book.remove(book.active)
    for name in ("Roster_Control", "Session_Control", "Reminder_Control", "Delivery_Summary"):
        sheet = book.create_sheet(name); sheet.append(["Incomplete"])
    book.save(root / REQUIRED[0]); book.close()
    document = Document(); document.add_paragraph("Team Training Coordination Brief"); document.add_paragraph("Package is incomplete."); document.save(root / REQUIRED[1])
    (root / REQUIRED[2]).write_bytes(b"BEGIN:VCALENDAR\r\nVERSION:2.0\r\nMETHOD:PUBLISH\r\nBEGIN:VEVENT\r\nUID:incomplete@evocowork.invalid\r\nDTSTART:20200508T090000\r\nEND:VEVENT\r\nEND:VCALENDAR\r\n")
    message = EmailMessage(); message["From"] = "Training Coordinator <training-coordinator@example.invalid>"; message["To"] = "Alice <alice@example.invalid>"; message["Subject"] = "Team training reminder"; message.set_content("Incomplete controlled draft; has not been sent.")
    raw_message = tempfile.NamedTemporaryFile(delete=False); raw_message.close()
    try:
        with open(raw_message.name, "wb") as handle:
            BytesGenerator(handle, policy=policy.SMTP).flatten(message)
        (root / REQUIRED[3]).write_bytes(b"From training-coordinator@example.invalid Fri May 01 10:00:00 2020\n" + Path(raw_message.name).read_bytes() + b"\n")
    finally:
        Path(raw_message.name).unlink(missing_ok=True)


def _input_only(root: Path) -> None:
    for path in list(root.iterdir()):
        if path.is_file():
            path.unlink()
    for source in INPUTS.iterdir():
        if source.is_file():
            shutil.copy2(source, root / source.name)


PARTIALS: dict[str, tuple[Mutator, bool]] = {
    "roster_name_wrong": (_excel("Roster_Control", "B2", "Alicia"), False),
    "roster_section_wrong": (_excel("Roster_Control", "C2", "Section C"), False),
    "roster_source_time_wrong": (_excel("Roster_Control", "D2", "10:00 AM - 12:00 PM"), False),
    "roster_training_date_wrong": (_excel("Roster_Control", "H2", "2020-05-15"), False),
    "roster_duration_formula_flattened": (_excel("Roster_Control", "G2", 2), True),
    "roster_address_formula_wrong": (_excel("Roster_Control", "I2", "alice@example.com"), True),
    "session_start_wrong": (_excel("Session_Control", "C2", "10:00"), False),
    "session_member_list_wrong": (_excel("Session_Control", "G2", "Alice | David | Grace"), False),
    "session_count_formula_flattened": (_excel("Session_Control", "F2", 4), True),
    "session_gap_formula_wrong": (_excel("Session_Control", "I4", "=0"), True),
    "summary_formula_flattened": (_excel("Delivery_Summary", "B4", 12), True),
    "reminder_recipient_wrong": (_excel("Reminder_Control", "D2", "alice@example.com"), True),
    "reminder_subject_wrong": (_excel("Reminder_Control", "E2", "Team meeting"), True),
    "native_chart_removed": (_remove_chart, False),
    "native_freeze_removed": (_remove_freeze, False),
    "master_drop_event": (_drop_master_event, False),
    "master_wrong_start": (_ics_replace(b"DTSTART:20200508T090000", b"DTSTART:20200508T100000"), False),
    "master_wrong_uid": (_ics_replace(b"UID:section-a-training-20200508@evocowork.invalid", b"UID:wrong-a@evocowork.invalid"), False),
    "master_wrong_roster": (_ics_replace(b"X-EVOCOWORK-ROSTER:Alice | David | Grace | Jack", b"X-EVOCOWORK-ROSTER:Alice | David | Grace"), False),
    "master_timezone_invented": (_ics_replace(b"DTSTART:20200508T090000", b"DTSTART;TZID=America/New_York:20200508T090000"), False),
    "master_location_invented": (_ics_replace(b"STATUS:TENTATIVE\r\n", b"STATUS:TENTATIVE\r\nLOCATION:Room 7\r\n"), False),
    "master_status_wrong": (_ics_replace(b"STATUS:TENTATIVE", b"STATUS:CONFIRMED"), False),
    "mbox_drop_message": (_mail(_drop_message), False),
    "mbox_duplicate_recipient": (_mail(_duplicate_recipient), False),
    "mbox_wrong_subject": (_replace_header(0, "Subject", "Team meeting"), False),
    "mbox_wrong_session_uid": (_replace_header(0, "X-EvoCowork-Session-UID", "wrong@evocowork.invalid"), False),
    "mbox_missing_body_fact": (_body_replace(0, "Location: not provided", "Location pending"), False),
    "mbox_drop_attachment": (_mail(_drop_attachment), False),
    "mbox_attachment_time_wrong": (_mail(_attachment_time), False),
    "mbox_real_recipient": (_replace_header(0, "To", "Alice <alice@company.com>"), False),
    "mbox_sent_claim": (_body_replace(0, "has not been sent", "has been sent"), False),
    "mbox_extra_cc": (_mail(_add_cc), False),
    "word_member_wrong": (_word_replace("Alice", "Alicia"), False),
    "word_time_wrong": (_word_replace("09:00-11:00", "10:00-12:00"), False),
    "word_gap_missing": (_word_replace("60-minute inter-session gap", "inter-session gap"), False),
    "word_sent_claim": (_word_replace("not sent communications", "sent communications"), False),
    "word_real_email": (_word_replace("alice@example.invalid", "alice@company.com"), False),
    "word_table_removed": (_word_remove_table, False),
    "unexpected_debug_file": (lambda root: (root / "debug.txt").write_text("temporary\n", encoding="utf-8"), False),
    "moderate_communication_domain_conflict": (_moderate_communication_domain, False),
    "severe_valid_but_incomplete_package": (_severe_valid_incomplete, False),
}


def main() -> None:
    parser = argparse.ArgumentParser(); parser.add_argument("--report", type=Path); args = parser.parse_args()
    golden_recalc, _ = JUDGE._recalculate(GOLDEN / REQUIRED[0]); golden_renders = {name: JUDGE._render(GOLDEN / name) for name in REQUIRED[:2]}
    original_recalc, original_render = JUDGE._recalculate, JUDGE._render
    assert JUDGE.evaluate(GOLDEN)["reward"] == 1.0

    def fast_recalc(path: Path):
        book = load_workbook(path, data_only=False); values = {}
        for sheet in book.worksheets:
            rows = []
            for row_index, row in enumerate(sheet.iter_rows(), 1):
                current = []
                for column_index, cell in enumerate(row, 1):
                    value = cell.value
                    if isinstance(value, str) and value.startswith("="):
                        value = JUDGE._value(golden_recalc, sheet.title, row_index, column_index)
                    current.append(value)
                rows.append(current)
            values[sheet.title] = rows
        book.close(); return values, "task-specific cached Golden formula values"

    def evaluate_case(case: Path, real_recalc: bool) -> float:
        JUDGE._recalculate = original_recalc if real_recalc else fast_recalc
        JUDGE._render = lambda path: golden_renders.get(path.name, (True, 1, "cached task render"))
        try:
            return float(JUDGE.evaluate(case)["reward"])
        finally:
            JUDGE._recalculate, JUDGE._render = original_recalc, original_render

    records: list[dict[str, object]] = []
    with tempfile.TemporaryDirectory(prefix="training-v2-negatives-") as raw:
        base = Path(raw)
        for name in REQUIRED:
            for mode in ("missing", "empty", "corrupt"):
                case = base / f"fatal-{mode}-{name}"; _copy(case); target = case / name
                if mode == "missing": target.unlink()
                elif mode == "empty": target.write_bytes(b"")
                else: target.write_bytes(b"not-a-valid-task-artifact")
                reward = evaluate_case(case, False); assert reward == 0, (mode, name, reward)
                records.append({"case": f"{mode}:{name}", "layer": "fatal_zero", "reward": reward})
        empty = base / "fatal-empty"; empty.mkdir(); reward = evaluate_case(empty, False); assert reward == 0; records.append({"case": "empty-output", "layer": "fatal_zero", "reward": reward})
        input_only = base / "fatal-input-only"; _copy(input_only); _input_only(input_only); reward = evaluate_case(input_only, False); assert reward == 0; records.append({"case": "unmodified-input-only", "layer": "fatal_zero", "reward": reward})
        for name, (mutator, real_recalc) in PARTIALS.items():
            case = base / f"partial-{name}"; _copy(case); mutator(case); reward = evaluate_case(case, real_recalc)
            if name in {"moderate_communication_domain_conflict", "severe_valid_but_incomplete_package"}:
                reward = float(JUDGE.evaluate(case)["reward"])
            assert 0 < reward < 1, (name, reward)
            records.append({"case": name, "layer": "partial", "reward": reward})
    partial = [float(record["reward"]) for record in records if record["layer"] == "partial"]
    by_name = {str(record["case"]): float(record["reward"]) for record in records}
    assert by_name["moderate_communication_domain_conflict"] < 0.8
    assert 0 < by_name["severe_valid_but_incomplete_package"] < by_name["moderate_communication_domain_conflict"]
    report = {"total": len(records), "fatal_zero": sum(record["layer"] == "fatal_zero" for record in records), "partial_continuous": len(partial), "partial_reward_range": [min(partial), max(partial)], "moderate_communication_domain_conflict": by_name["moderate_communication_domain_conflict"], "severe_valid_but_incomplete": by_name["severe_valid_but_incomplete_package"], "cases": records}
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True); args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
