#!/usr/bin/env python3
"""构造并回放 OfficeBench 3-6 团队培训任务的定向负例。"""

from __future__ import annotations

import argparse
import mailbox
import shutil
import tempfile
from email import policy
from email.generator import BytesGenerator
from email.parser import BytesParser
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from judge import DOCX_NAME, ICS_NAME, MBOX_NAME, REQUIRED_FILES, XLSX_NAME, evaluate


CASES = (
    "missing-docx",
    "formulas-flattened",
    "roster-name-changed",
    "roster-section-changed",
    "session-time-changed",
    "session-date-changed",
    "session-gap-formula-changed",
    "summary-staticized",
    "native-chart-removed",
    "master-drop-event",
    "master-timezone-invented",
    "master-location-added",
    "mbox-drop-message",
    "mbox-duplicate-recipient",
    "mbox-wrong-section",
    "mbox-real-recipient",
    "mbox-drop-attachment",
    "mbox-attachment-time-changed",
    "word-sent-claim",
    "word-roster-conflict",
    "extra-debug-file",
    "corrupt-xlsx",
)


def copy_golden(golden: Path, case_dir: Path) -> None:
    case_dir.mkdir(parents=True, exist_ok=True)
    for filename in REQUIRED_FILES:
        shutil.copy2(golden / filename, case_dir / filename)


def replace_docx_text(path: Path, old: str, new: str) -> None:
    document = Document(path)
    locations = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                locations.extend(cell.paragraphs)
    replaced = False
    for paragraph in locations:
        if old not in paragraph.text:
            continue
        replaced = True
        if any(old in run.text for run in paragraph.runs):
            for run in paragraph.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
        else:
            paragraph.text = paragraph.text.replace(old, new)
    if not replaced:
        raise ValueError(f"未找到 Word 待替换文本：{old}")
    document.save(path)


def message_factory(handle):
    return BytesParser(policy=policy.default).parse(handle)


def read_messages(path: Path):
    return list(mailbox.mbox(path, factory=message_factory, create=False))


def write_messages(path: Path, messages) -> None:
    with tempfile.TemporaryDirectory(prefix="training-negative-mbox-") as temporary:
        target = Path(temporary) / "mutated.mbox"
        box = mailbox.mbox(target, create=True)
        try:
            box.lock()
            for message in messages:
                if message.get_unixfrom() is None:
                    message.set_unixfrom("From training-coordinator@example.invalid Fri May 01 10:00:00 2020")
                box.add(message)
            box.flush()
        finally:
            try:
                box.unlock()
            except Exception:
                pass
            box.close()
        shutil.copy2(target, path)


def mutate(case: str, case_dir: Path) -> None:
    workbook_path = case_dir / XLSX_NAME
    document_path = case_dir / DOCX_NAME
    calendar_path = case_dir / ICS_NAME
    mbox_path = case_dir / MBOX_NAME
    if case == "missing-docx":
        document_path.unlink()
    elif case == "formulas-flattened":
        formulas = load_workbook(workbook_path, data_only=False)
        values = load_workbook(workbook_path, data_only=True)
        for formula_sheet, value_sheet in zip(formulas.worksheets, values.worksheets, strict=True):
            for row in formula_sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        cell.value = value_sheet[cell.coordinate].value
        formulas.save(workbook_path)
    elif case == "roster-name-changed":
        workbook = load_workbook(workbook_path)
        workbook["Roster_Control"]["B2"] = "Alicia"
        workbook.save(workbook_path)
    elif case == "roster-section-changed":
        workbook = load_workbook(workbook_path)
        workbook["Roster_Control"]["C2"] = "Section C"
        workbook.save(workbook_path)
    elif case == "session-time-changed":
        workbook = load_workbook(workbook_path)
        workbook["Session_Control"]["C2"] = "10:00"
        workbook.save(workbook_path)
    elif case == "session-date-changed":
        workbook = load_workbook(workbook_path)
        workbook["Session_Control"]["B2"] = "2020-05-15"
        workbook.save(workbook_path)
    elif case == "session-gap-formula-changed":
        workbook = load_workbook(workbook_path)
        workbook["Session_Control"]["I4"] = "=0"
        workbook.save(workbook_path)
    elif case == "summary-staticized":
        workbook = load_workbook(workbook_path)
        workbook["Delivery_Summary"]["B4"] = 12
        workbook.save(workbook_path)
    elif case == "native-chart-removed":
        workbook = load_workbook(workbook_path)
        workbook["Delivery_Summary"]._charts = []
        workbook.save(workbook_path)
    elif case == "master-drop-event":
        raw = calendar_path.read_bytes()
        start = raw.find(b"BEGIN:VEVENT\r\n", raw.find(b"BEGIN:VEVENT\r\n") + 1)
        end = raw.find(b"END:VEVENT\r\n", start) + len(b"END:VEVENT\r\n")
        calendar_path.write_bytes(raw[:start] + raw[end:])
    elif case == "master-timezone-invented":
        raw = calendar_path.read_bytes()
        raw = raw.replace(b"DTSTART:20200508T090000", b"DTSTART;TZID=America/New_York:20200508T090000", 1)
        calendar_path.write_bytes(raw)
    elif case == "master-location-added":
        raw = calendar_path.read_bytes()
        raw = raw.replace(b"STATUS:TENTATIVE\r\n", b"STATUS:TENTATIVE\r\nLOCATION:Room 7\r\n", 1)
        calendar_path.write_bytes(raw)
    elif case == "mbox-drop-message":
        messages = read_messages(mbox_path)
        write_messages(mbox_path, messages[:-1])
    elif case == "mbox-duplicate-recipient":
        messages = read_messages(mbox_path)
        messages[1].replace_header("To", "Alice <alice@example.invalid>")
        write_messages(mbox_path, messages)
    elif case == "mbox-wrong-section":
        messages = read_messages(mbox_path)
        messages[0].replace_header("Subject", "Team training reminder - Section C - 2020-05-08")
        write_messages(mbox_path, messages)
    elif case == "mbox-real-recipient":
        messages = read_messages(mbox_path)
        messages[0].replace_header("To", "Alice <alice@example.com>")
        write_messages(mbox_path, messages)
    elif case == "mbox-drop-attachment":
        messages = read_messages(mbox_path)
        payload = messages[0].get_payload()
        if not isinstance(payload, list) or len(payload) < 2:
            raise ValueError("首封邮件没有可删除的附件")
        messages[0].set_payload(payload[:-1])
        write_messages(mbox_path, messages)
    elif case == "mbox-attachment-time-changed":
        messages = read_messages(mbox_path)
        attachment = next(messages[0].iter_attachments())
        raw = (attachment.get_payload(decode=True) or b"").replace(
            b"DTSTART:20200508T090000", b"DTSTART:20200508T100000", 1
        )
        filename = attachment.get_filename()
        attachment.set_content(
            raw,
            maintype="text",
            subtype="calendar",
            cte="base64",
            params={"method": "PUBLISH", "charset": "utf-8"},
        )
        attachment.add_header("Content-Disposition", "attachment", filename=filename)
        write_messages(mbox_path, messages)
    elif case == "word-sent-claim":
        document = Document(document_path)
        document.add_paragraph("All 12 reminder emails have been sent successfully.")
        document.save(document_path)
    elif case == "word-roster-conflict":
        replace_docx_text(document_path, "Alice", "Alicia")
    elif case == "extra-debug-file":
        (case_dir / "debug.txt").write_text("temporary", encoding="utf-8")
    elif case == "corrupt-xlsx":
        workbook_path.write_bytes(b"not-an-xlsx")
    else:
        raise ValueError(case)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--case", choices=CASES, action="append", required=True)
    args = parser.parse_args()
    failures: list[dict[str, object]] = []
    for case in args.case:
        with tempfile.TemporaryDirectory(prefix=f"training-negative-{case}-") as temporary:
            case_dir = Path(temporary)
            copy_golden(args.golden, case_dir)
            mutate(case, case_dir)
            result = evaluate(case_dir)
            print(
                f"{case}\treward={result['reward']:.6f}\tpassed={result['passed']}"
                f"\thurdles={','.join(result['hurdle_failures'])}",
                flush=True,
            )
            if result["passed"] or float(result["reward"]) >= 0.8:
                failures.append({"case": case, "result": result})
    if failures:
        raise SystemExit(f"负例错误通过：{[item['case'] for item in failures]}")


if __name__ == "__main__":
    main()
