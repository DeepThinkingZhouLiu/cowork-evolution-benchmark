#!/usr/bin/env python3
"""根据逐项目视抽取的 OfficeBench 3-69 成绩单真值构建 Golden。"""

from __future__ import annotations

import argparse
import hashlib
import re
import shutil
import subprocess
import tempfile
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.table import Table, TableStyleInfo


XLSX_NAME = "Transcript_GPA_Audit.xlsx"
DOCX_NAME = "Transcript_GPA_Audit_Report.docx"
SOURCE_NAME = "transcripts.pdf"
SOURCE_HASH = "e471b9567cbce1b0335d6d097bc068d4b61d0223d49249ad5dc87095d5d80318"
FIXED_DT = datetime(2026, 9, 1, 12, 5, 0, tzinfo=timezone.utc)

# 这些记录来自对图像型 PDF 的逐行人工目视抽取，不是随机生成数据。
COURSES = [
    ("2014/2015", "Fall", "Physical Science", "", "A", 1.0),
    ("2014/2015", "Fall", "U.S. History", "", "A", 1.0),
    ("2014/2015", "Winter", "Tennis", "", "A", 1.0),
    ("2014/2015", "Spring", "English Composition", "", "A", 1.0),
    ("2014/2015", "Spring", "Music Appreciation", "", "B", 1.0),
    ("2014/2015", "Summer", "Algebra I", "", "A", 1.0),
    ("2015/2016", "Fall", "American Literature", "", "A", 1.0),
    ("2015/2016", "Fall", "Art Appreciation", "", "A", 1.0),
    ("2015/2016", "Winter", "Soccer", "", "A", 1.0),
    ("2015/2016", "Spring", "Algebra II", "", "B", 1.0),
    ("2015/2016", "Spring", "Biology", "", "A", 1.0),
    ("2015/2016", "Summer", "World History", "", "A", 1.0),
    ("2016/2017", "Fall", "Modern Literature", "", "B", 1.0),
    ("2016/2017", "Fall", "Music Instruction", "", "C", 1.0),
    ("2016/2017", "Winter", "Baseball", "", "A", 1.0),
    ("2016/2017", "Spring", "Geometry", "", "A", 1.0),
    ("2016/2017", "Spring", "Chemistry", "", "C", 1.0),
    ("2016/2017", "Summer", "Trigonometry", "", "A", 1.0),
    ("2017/2018", "Fall", "Geography", "", "A", 1.0),
    ("2017/2018", "Fall", "Physics", "", "A", 1.0),
    ("2017/2018", "Winter", "Economics", "", "A", 1.0),
    ("2017/2018", "Spring", "Drama", "DE", "A", 1.0),
    ("2017/2018", "Spring", "World Literature", "AP", "B", 1.0),
]

# PDF 每学期直接显示的 GPA、累计 GPA 和累计学分。
TERMS = [
    ("2014/2015", "Fall", 4.00, 4.00, 2.0),
    ("2014/2015", "Winter", 4.00, 4.00, 3.0),
    ("2014/2015", "Spring", 3.50, 3.80, 5.0),
    ("2014/2015", "Summer", 4.00, 3.83, 6.0),
    ("2015/2016", "Fall", 4.00, 3.88, 8.0),
    ("2015/2016", "Winter", 4.00, 3.89, 9.0),
    ("2015/2016", "Spring", 3.50, 3.82, 11.0),
    ("2015/2016", "Summer", 4.00, 3.83, 12.0),
    ("2016/2017", "Fall", 2.50, 3.64, 14.0),
    ("2016/2017", "Winter", 4.00, 3.67, 15.0),
    ("2016/2017", "Spring", 3.00, 3.59, 17.0),
    ("2016/2017", "Summer", 4.00, 3.61, 18.0),
    ("2017/2018", "Fall", 4.00, 3.65, 20.0),
    ("2017/2018", "Winter", 4.00, 3.67, 21.0),
    ("2017/2018", "Spring", 5.00, 3.78, 23.0),
]

NAVY = "17324D"
BLUE = "276B8E"
TEAL = "168AAD"
AMBER = "D69E2E"
RED = "B5403C"
WHITE = "FFFFFF"
PALE_BLUE = "EAF3F8"
PALE_GREEN = "E7F4EC"
PALE_AMBER = "FFF4D6"
PALE_RED = "FBE9E7"
LINE = "CAD6E0"
GRAY = "607080"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def validate_source(input_dir: Path) -> Path:
    source = input_dir / SOURCE_NAME
    if not source.is_file():
        raise FileNotFoundError(source)
    actual_hash = sha256(source)
    if actual_hash != SOURCE_HASH:
        raise ValueError(f"输入 SHA256 不一致：{actual_hash}")
    proc = subprocess.run(["pdfinfo", str(source)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=60)
    match = re.search(r"^Pages:\s+(\d+)", proc.stdout, flags=re.MULTILINE)
    if proc.returncode != 0 or not match or int(match.group(1)) != 1:
        raise ValueError("冻结输入必须是实际只有 1 页的 PDF")
    return source


def grade_points(grade: str) -> int:
    return {"A": 4, "B": 3, "C": 2, "D": 1, "F": 0}[grade]


def assert_truth() -> None:
    if len(COURSES) != 23 or len(TERMS) != 15:
        raise ValueError("人工抽取数量不一致")
    grades = Counter(row[4] for row in COURSES)
    quality = sum(grade_points(row[4]) * row[5] for row in COURSES)
    if grades != Counter({"A": 17, "B": 4, "C": 2}) or quality != 84 or sum(row[5] for row in COURSES) != 23:
        raise ValueError("成绩单真值自检失败")
    if TERMS[-1][2:] != (5.0, 3.78, 23.0):
        raise ValueError("最后学期来源值不一致")


def style_sheet(sheet, widths: dict[str, float], orientation: str = "landscape") -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 34
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color=LINE)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=9, color=NAVY)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.page_setup.orientation = orientation
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:1"
    sheet.page_margins.left = 0.2
    sheet.page_margins.right = 0.2
    sheet.page_margins.top = 0.3
    sheet.page_margins.bottom = 0.3


def add_table(sheet, name: str, ref: str | None = None) -> None:
    table = Table(displayName=name, ref=ref or sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    sheet.add_table(table)


def build_workbook(output_path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.properties.creator = "EvoCowork OfficeBench adapter"
    workbook.properties.title = "Transcript GPA Audit"
    workbook.properties.created = FIXED_DT.replace(tzinfo=None)
    workbook.properties.modified = FIXED_DT.replace(tzinfo=None)

    profile = workbook.create_sheet("Source_Profile")
    profile.append(["Field", "Extracted_Value", "Evidence_Location", "Audit_Note"])
    profile_rows = [
        ("Student", "Teddy Roosevelt", "Header — Student", "Academic identifier retained; direct contact details excluded"),
        ("Issue_Date", "October 25, 2018", "Header — Date Issued", "Source-reported date"),
        ("Credits_Earned", 23.0, "Header — Credits Earned", "Matches visible course-credit total"),
        ("Weighted_GPA", 3.78, "Header — Weighted GPA", "Source-reported; weighting policy not supplied"),
        ("Unweighted_GPA", 3.65, "Header — Unweighted GPA", "Matches independent A/B/C recalculation after rounding"),
        ("PDF_File_Type", "Image-only", "PDF object inspection", "No extractable text layer"),
        ("Delivered_PDF_Pages", 1, "PDF page tree", "Only one page is present in the delivered file"),
        ("Printed_Page_Marker", "1 of 2", "Visible page footer", "Printed marker claims a two-page source"),
        ("Printed_Total_Pages", 2, "Visible page footer", "Do not infer missing-page content"),
        ("Source_Completeness_Risk", "Potentially incomplete", "Delivered pages vs printed marker", "Escalate before a completeness-sensitive decision"),
    ]
    for row in profile_rows:
        profile.append(row)
    style_sheet(profile, {"A": 28, "B": 24, "C": 29, "D": 55}, orientation="landscape")
    for row in range(2, 12):
        profile.cell(row, 1).fill = PatternFill("solid", fgColor=PALE_BLUE)
        profile.cell(row, 1).font = Font(name="Aptos", size=9, bold=True, color=NAVY)
    for cell in (profile["B4"], profile["B5"], profile["B6"]):
        cell.number_format = "0.00"
    profile.conditional_formatting.add("B11", FormulaRule(formula=['B11="Potentially incomplete"'], fill=PatternFill("solid", fgColor=PALE_AMBER)))
    add_table(profile, "TranscriptSourceProfileTable")
    profile.print_area = "A1:D11"

    ledger = workbook.create_sheet("Course_Ledger")
    ledger.append([
        "Sequence", "Academic_Year", "Term", "Course", "Program_Tag", "Grade", "Credits",
        "Unweighted_Grade_Points", "Unweighted_Quality_Points",
    ])
    for index, (year, term, course, tag, grade, credits) in enumerate(COURSES, start=1):
        row = index + 1
        ledger.append([
            index, year, term, course, tag, grade, credits,
            f'=IF(F{row}="A",4,IF(F{row}="B",3,IF(F{row}="C",2,IF(F{row}="D",1,IF(F{row}="F",0,"")))))',
            f"=G{row}*H{row}",
        ])
    style_sheet(ledger, {"A": 10, "B": 15, "C": 12, "D": 25, "E": 12, "F": 9, "G": 10, "H": 24, "I": 26})
    for row in range(2, 25):
        ledger.cell(row, 7).number_format = "0.00"
        ledger.cell(row, 8).number_format = "0.00"
        ledger.cell(row, 9).number_format = "0.00"
    ledger.conditional_formatting.add("E2:E24", FormulaRule(formula=['OR(E2="DE",E2="AP")'], fill=PatternFill("solid", fgColor=PALE_AMBER)))
    ledger.conditional_formatting.add("F2:F24", FormulaRule(formula=['F2="C"'], fill=PatternFill("solid", fgColor=PALE_RED)))
    add_table(ledger, "TranscriptCourseLedgerTable")
    ledger.print_area = "A1:I24"

    term_audit = workbook.create_sheet("Term_Audit")
    term_audit.append([
        "Academic_Year", "Term", "Course_Count", "Credits", "Unweighted_Quality_Points",
        "Calculated_Unweighted_Term_GPA", "Displayed_Term_GPA",
    ])
    for index, (year, term, displayed_term, displayed_cumulative, cumulative_credits) in enumerate(TERMS, start=2):
        term_audit.append([
            year,
            term,
            f'=COUNTIFS(Course_Ledger!$B$2:$B$24,A{index},Course_Ledger!$C$2:$C$24,B{index})',
            f'=SUMIFS(Course_Ledger!$G$2:$G$24,Course_Ledger!$B$2:$B$24,A{index},Course_Ledger!$C$2:$C$24,B{index})',
            f'=SUMIFS(Course_Ledger!$I$2:$I$24,Course_Ledger!$B$2:$B$24,A{index},Course_Ledger!$C$2:$C$24,B{index})',
            f'=IFERROR(E{index}/D{index},"")',
            displayed_term,
        ])
        if abs(cumulative_credits - sum(row[5] for row in COURSES[: sum(1 for course in COURSES if (course[0], course[1]) in [(item[0], item[1]) for item in TERMS[: index - 1]])])) > 1e-9:
            raise ValueError("学期累计学分自检失败")
    term_audit.append([])
    term_audit.append([
        "Academic_Year", "Term", "Cumulative_Credits", "Cumulative_Unweighted_Points",
        "Calculated_Cumulative_Unweighted_GPA", "Displayed_Cumulative_GPA", "Displayed_vs_Calculated_Delta",
    ])
    cumulative_header_row = 18
    for offset, (year, term, _, displayed_cumulative, _) in enumerate(TERMS, start=1):
        row = cumulative_header_row + offset
        detail_row = offset + 1
        term_audit.append([
            year,
            term,
            f"=SUM($D$2:D{detail_row})",
            f"=SUM($E$2:E{detail_row})",
            f'=IFERROR(D{row}/C{row},"")',
            displayed_cumulative,
            f"=F{row}-E{row}",
        ])
    style_sheet(term_audit, {"A": 16, "B": 12, "C": 21, "D": 30, "E": 33, "F": 27, "G": 31})
    term_audit.page_setup.fitToHeight = 2
    term_audit.print_title_rows = None
    # 两块原生表纵向分页，避免横向分页丢失学年/学期上下文。
    term_audit.row_breaks.append(Break(id=17))
    for cell in term_audit[cumulative_header_row]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    term_audit.row_dimensions[cumulative_header_row].height = 34
    for row in range(2, 17):
        for column in (4, 5, 6, 7):
            term_audit.cell(row, column).number_format = "0.0000"
    for row in range(19, 34):
        for column in (3, 4, 5, 6, 7):
            term_audit.cell(row, column).number_format = "0.0000"
    term_audit.conditional_formatting.add("G19:G33", FormulaRule(formula=["ABS(G19)>0.02"], fill=PatternFill("solid", fgColor=PALE_AMBER)))
    add_table(term_audit, "TranscriptTermDetailTable", "A1:G16")
    add_table(term_audit, "TranscriptCumulativeAuditTable", "A18:G33")
    term_audit.print_area = "A1:G33"

    summary = workbook.create_sheet("Summary")
    summary.append(["Metric", "Value", "Interpretation"])
    metrics = [
        ("Extracted_Courses", "=COUNTA(Course_Ledger!A2:A24)", "Visible courses only"),
        ("Extracted_Credits", "=SUM(Course_Ledger!G2:G24)", "Must reconcile to header credits"),
        ("Grade_A_Count", '=COUNTIF(Course_Ledger!F2:F24,"A")', "Unweighted grade distribution"),
        ("Grade_B_Count", '=COUNTIF(Course_Ledger!F2:F24,"B")', "Unweighted grade distribution"),
        ("Grade_C_Count", '=COUNTIF(Course_Ledger!F2:F24,"C")', "Unweighted grade distribution"),
        ("Unweighted_Quality_Points", "=SUM(Course_Ledger!I2:I24)", "A=4, B=3, C=2, D=1, F=0"),
        ("Calculated_Unweighted_GPA", "=B7/B3", "Independent course-level recalculation"),
        ("Rounded_Calculated_Unweighted_GPA", "=ROUND(B8,2)", "Two-decimal comparison value"),
        ("Reported_Unweighted_GPA", "=Source_Profile!B6", "Source header value"),
        ("Rounded_Unweighted_Delta", "=B9-B10", "Expected zero"),
        ("Reported_Weighted_GPA", "=Source_Profile!B5", "Source-reported; policy absent"),
        ("Final_Displayed_Cumulative_GPA", "=Term_Audit!F33", "Last visible term"),
        ("Weighted_vs_Final_Displayed_Delta", "=B13-B12", "Expected zero"),
        ("Delivered_PDF_Pages", "=Source_Profile!B8", "Actual PDF page tree"),
        ("Printed_Total_Pages", "=Source_Profile!B10", "Visible footer claim"),
        ("Source_Completeness_Flag", '=IF(B15<B16,"Potentially incomplete","Complete")', "Do not infer missing-page content"),
    ]
    for row in metrics:
        summary.append(row)
    style_sheet(summary, {"A": 39, "B": 24, "C": 48}, orientation="landscape")
    for row in range(2, 18):
        summary.cell(row, 1).fill = PatternFill("solid", fgColor=PALE_BLUE)
        summary.cell(row, 1).font = Font(name="Aptos", size=9, bold=True, color=NAVY)
        summary.cell(row, 2).font = Font(name="Aptos Display", size=11, bold=True, color=TEAL)
    for row in range(3, 16):
        summary.cell(row, 2).number_format = "0.0000"
    summary.conditional_formatting.add("B17", FormulaRule(formula=['B17="Potentially incomplete"'], fill=PatternFill("solid", fgColor=PALE_AMBER)))
    add_table(summary, "TranscriptAuditSummaryTable")
    chart = LineChart()
    chart.style = 13
    chart.title = "Cumulative GPA: recalculated vs displayed"
    chart.y_axis.title = "GPA"
    chart.x_axis.title = "Term sequence"
    chart.height = 7.2
    chart.width = 14.5
    chart.add_data(Reference(term_audit, min_col=5, max_col=6, min_row=18, max_row=33), titles_from_data=True)
    chart.series[0].tx = SeriesLabel(v="Calculated unweighted")
    chart.series[1].tx = SeriesLabel(v="Displayed")
    chart.set_categories(Reference(term_audit, min_col=2, min_row=19, max_row=33))
    chart.y_axis.scaling.min = 3.0
    chart.y_axis.scaling.max = 4.1
    summary.add_chart(chart, "E2")
    summary.print_area = "A1:O17"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    recalculate_workbook(output_path)
    restore_frozen_rows(output_path, 4)
    normalize_ooxml(output_path)


def recalculate_workbook(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="transcript-xlsx-recalc-") as tmp:
        root = Path(tmp)
        input_root = root / "input"
        output_root = root / "output"
        profile = root / "profile"
        input_root.mkdir()
        output_root.mkdir()
        profile.mkdir()
        source = input_root / path.name
        shutil.copy2(path, source)
        proc = subprocess.run(
            ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "xlsx", "--outdir", str(output_root), str(source)],
            stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300,
        )
        converted = output_root / path.name
        if proc.returncode != 0 or not converted.is_file():
            raise RuntimeError(f"LibreOffice 公式重算失败：{proc.stdout[-1000:]}")
        shutil.copy2(converted, path)


def restore_frozen_rows(path: Path, sheet_count: int) -> None:
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with tempfile.TemporaryDirectory(prefix="transcript-xlsx-freeze-") as tmp:
        target = Path(tmp) / path.name
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w") as destination:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename in {f"xl/worksheets/sheet{index}.xml" for index in range(1, sheet_count + 1)}:
                    root = etree.fromstring(data)
                    sheet_view = root.find(f"{{{namespace}}}sheetViews/{{{namespace}}}sheetView")
                    if sheet_view is not None:
                        for existing in sheet_view.findall(f"{{{namespace}}}pane"):
                            sheet_view.remove(existing)
                        pane = etree.Element(f"{{{namespace}}}pane")
                        pane.set("ySplit", "1")
                        pane.set("topLeftCell", "A2")
                        pane.set("activePane", "bottomLeft")
                        pane.set("state", "frozen")
                        sheet_view.insert(0, pane)
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                destination.writestr(info, data)
        shutil.copy2(target, path)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE, size: str = "6") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def add_word_table(document: Document, headers: list[str], data: list[list[str]], widths: list[float], font_size: float = 8.0):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, NAVY)
        set_cell_border(cell, NAVY)
        cell.width = Cm(widths[index])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(font_size)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(data):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.text = str(value)
            set_cell_shading(cell, WHITE if row_index % 2 == 0 else "F4F7FA")
            set_cell_border(cell)
            cell.width = Cm(widths[index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = RGBColor.from_string(NAVY)
    return table


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Transcript audit  |  ")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def build_report(output_path: Path) -> None:
    grades = Counter(row[4] for row in COURSES)
    quality = sum(grade_points(row[4]) * row[5] for row in COURSES)
    calculated = quality / sum(row[5] for row in COURSES)
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    document.core_properties.creator = "EvoCowork OfficeBench adapter"
    document.core_properties.title = "Transcript GPA Audit Report"
    document.core_properties.created = FIXED_DT.replace(tzinfo=None)
    document.core_properties.modified = FIXED_DT.replace(tzinfo=None)
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(9.5)
    document.styles["Normal"].font.color.rgb = RGBColor.from_string(NAVY)
    for name in ("Title", "Heading 1", "Heading 2"):
        document.styles[name].font.name = "Aptos Display"
        document.styles[name].font.color.rgb = RGBColor.from_string(NAVY)

    title = document.add_paragraph(style="Title")
    title_run = title.add_run("Transcript GPA Audit Report")
    title_run.font.size = Pt(26)
    title_run.font.bold = True
    subtitle = document.add_paragraph()
    subtitle_run = subtitle.add_run("Teddy Roosevelt  |  Issued October 25, 2018  |  Evidence-bounded review")
    subtitle_run.font.name = "Aptos"
    subtitle_run.font.size = Pt(10)
    subtitle_run.font.color.rgb = RGBColor.from_string(TEAL)

    add_word_table(
        document,
        ["Courses", "Credits", "Reported unweighted", "Reported weighted"],
        [["23", "23.00", "3.65 / 4.00", "3.78 / 4.00"]],
        [4.4, 4.4, 4.8, 4.8],
    )

    document.add_heading("Source Completeness", level=1)
    paragraph = document.add_paragraph()
    lead = paragraph.add_run("Completeness risk — ")
    lead.bold = True
    lead.font.color.rgb = RGBColor.from_string(RED)
    paragraph.add_run(
        "the delivered image-only PDF contains 1 actual page, while its visible footer says “1 of 2.” "
        "The file is therefore potentially incomplete. No content from the unprovided page is inferred."
    )

    document.add_heading("Reported Academic Profile", level=1)
    add_word_table(
        document,
        ["Field", "Source-reported value", "Audit treatment"],
        [
            ["Student", "Teddy Roosevelt", "Retained as the academic identifier"],
            ["Credits earned", "23.00", "Reconciled to the visible course ledger"],
            ["Unweighted GPA", "3.65 / 4.00", "Independently recalculated"],
            ["Weighted GPA", "3.78 / 4.00", "Retained as reported; policy unavailable"],
        ],
        [4.4, 5.1, 9.0],
    )

    document.add_heading("Course Extraction and Recalculation", level=1)
    document.add_paragraph(
        f"The visible page contains 23 courses across 15 terms and 23.00 credits. The grade distribution is "
        f"17 A, 4 B and 2 C. Using the stated audit scale A=4, B=3, C=2, D=1 and F=0 gives "
        f"84.00 unweighted quality points and {quality:.0f}/23 = {calculated:.9f}; rounded to two decimals, this is 3.65 and matches the reported unweighted GPA."
    )

    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_heading("Term-Level Reconciliation", level=1)
    term_rows = []
    cumulative_points = 0.0
    cumulative_credits = 0.0
    for year, term, displayed_term, displayed_cumulative, _ in TERMS:
        selected = [row for row in COURSES if row[0] == year and row[1] == term]
        credits = sum(row[5] for row in selected)
        points = sum(grade_points(row[4]) * row[5] for row in selected)
        cumulative_credits += credits
        cumulative_points += points
        term_rows.append([
            f"{year} {term}",
            str(len(selected)),
            f"{credits:.2f}",
            f"{points / credits:.2f}",
            f"{displayed_term:.2f}",
            f"{cumulative_points / cumulative_credits:.2f}",
            f"{displayed_cumulative:.2f}",
        ])
    add_word_table(
        document,
        ["Period", "Courses", "Credits", "Calc term", "Displayed term", "Calc cumulative", "Displayed cumulative"],
        term_rows,
        [4.0, 1.8, 1.8, 2.2, 2.4, 2.7, 2.7],
        font_size=6.8,
    )
    document.add_paragraph(
        "The final visible term, 2017/2018 Spring, contains Drama (DE) grade A and World Literature (AP) grade B. "
        "Its unweighted recalculation is 3.50, while the PDF displays a term GPA of 5.00, a difference of 1.50."
    )

    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_heading("Weighted GPA Evidence Boundary", level=1)
    document.add_paragraph(
        "The source reports a weighted GPA of 3.78 / 4.00, and its final visible cumulative GPA is also 3.78. "
        "Those values are retained as source-reported facts. The delivered material does not include a weighting policy, "
        "and the printed page marker indicates an unprovided second page. Therefore this audit cannot independently reconstruct, "
        "validate or assign the weighting rule behind 3.78 or the final-term 5.00. It does not assume conventional DE/AP bonuses."
    )
    add_word_table(
        document,
        ["Check", "Result", "Boundary"],
        [
            ["Unweighted course-level GPA", "84.00 / 23.00 = 3.652173913 → 3.65", "Recomputed and reconciled"],
            ["Reported weighted GPA", "3.78", "Retained; not independently reconstructed"],
            ["Final displayed cumulative GPA", "3.78", "Matches the reported weighted value"],
            ["Source completeness", "1 delivered page vs printed 1 of 2", "Potentially incomplete"],
        ],
        [5.0, 7.0, 6.5],
    )

    document.add_heading("Audit Conclusion", level=1)
    document.add_paragraph(
        "The visible course ledger reconciles to 23.00 credits, and the independently calculated unweighted GPA rounds to the reported 3.65. "
        "The weighted GPA should remain reported-but-unverified until the missing page and an authoritative weighting policy are obtained. "
        "This conclusion addresses numeric reconciliation only; it is not an authenticity, graduation or eligibility determination."
    )

    add_page_number(document.sections[0].footer.paragraphs[0])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    normalize_ooxml(output_path)


def normalize_ooxml(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="transcript-ooxml-normalize-") as tmp:
        target = Path(tmp) / path.name
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w") as destination:
            for info in sorted(source.infolist(), key=lambda item: item.filename):
                data = source.read(info.filename)
                if info.filename == "docProps/core.xml":
                    root = etree.fromstring(data)
                    dcterms = "http://purl.org/dc/terms/"
                    for field in ("created", "modified"):
                        node = root.find(f"{{{dcterms}}}{field}")
                        if node is not None:
                            node.text = "2026-09-01T12:05:00Z"
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                elif info.filename.startswith("xl/charts/") and info.filename.endswith(".xml"):
                    root = etree.fromstring(data)
                    chart_ns = "http://schemas.openxmlformats.org/drawingml/2006/chart"
                    category_axis = root.find(f".//{{{chart_ns}}}catAx/{{{chart_ns}}}axId")
                    value_axis = root.find(f".//{{{chart_ns}}}valAx/{{{chart_ns}}}axId")
                    if category_axis is not None and value_axis is not None:
                        mapping = {category_axis.get("val", ""): "36917001", value_axis.get("val", ""): "36917002"}
                        for node in root.iter():
                            if etree.QName(node).localname in {"axId", "crossAx"} and node.get("val") in mapping:
                                node.set("val", mapping[node.get("val")])
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                copied = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
                copied.compress_type = zipfile.ZIP_DEFLATED
                copied.external_attr = info.external_attr
                copied.create_system = info.create_system
                copied.flag_bits = info.flag_bits
                destination.writestr(copied, data)
        shutil.copy2(target, path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    validate_source(args.input_dir)
    assert_truth()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for child in args.output_dir.iterdir():
        if child.is_file():
            child.unlink()
        elif child.is_dir():
            shutil.rmtree(child)
    build_workbook(args.output_dir / XLSX_NAME)
    build_report(args.output_dir / DOCX_NAME)


if __name__ == "__main__":
    main()
