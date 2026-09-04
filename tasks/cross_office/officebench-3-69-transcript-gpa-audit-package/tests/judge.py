#!/usr/bin/env python3
"""OfficeBench 3-69 成绩单审计包的 v2 原子连续评分。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from openpyxl import load_workbook


TASK_ID = "cowork-evo/officebench-3-69-transcript-gpa-audit-package"
RUBRIC_VERSION = "officebench-3-69-transcript-gpa-audit-v0.2"
JUDGE_VERSION = "transcript-gpa-audit-atomic-v2.0"
XLSX = "Transcript_GPA_Audit.xlsx"
DOCX = "Transcript_GPA_Audit_Report.docx"
REQUIRED = (XLSX, DOCX)
SOURCE_HASH = "e471b9567cbce1b0335d6d097bc068d4b61d0223d49249ad5dc87095d5d80318"
WEIGHTS = {"R001":6,"R002":9,"R003":12,"R004":7,"R005":10,"R006":12,"R007":9,"R008":5,"R009":7,"R010":8,"R011":6,"R012":5,"R013":4}
PENALTIES = {"P001":-12,"P002":-4}


def _load_v1():
    path=Path(__file__).with_name("judge_v1.py");spec=importlib.util.spec_from_file_location("transcript_judge_v1",path)
    if spec is None or spec.loader is None:raise RuntimeError("无法载入 v1 Judge")
    module=importlib.util.module_from_spec(spec);sys.modules[spec.name]=module;spec.loader.exec_module(module);return module


V1=_load_v1()
COURSES=V1.COURSES
TERMS=V1.TERMS


def _sha256(path:Path)->str:
    digest=hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda:handle.read(1024*1024),b""):digest.update(block)
    return digest.hexdigest()


def _task_root()->Path:return Path(__file__).resolve().parents[1]


def _source()->Path:
    harbor=Path("/app/input_files/transcripts.pdf")
    return harbor if harbor.is_file() else _task_root()/"data"/"input_files"/"transcripts.pdf"


def _norm(value:Any)->str:return re.sub(r"\s+"," ",str(value or "").replace("\u00a0"," ").replace("\v"," ")).strip().casefold()


def _same(actual:Any,expected:Any,tolerance:float=1e-7)->bool:
    if isinstance(expected,(int,float)) and not isinstance(expected,bool):
        try:return abs(float(actual)-float(expected))<=tolerance
        except (TypeError,ValueError):return False
    return _norm(actual)==_norm(expected)


def _ratio(items:Iterable[tuple[bool,str]])->tuple[float,list[str],int,int]:
    findings=list(items);earned=sum(bool(ok) for ok,_ in findings)
    return earned/max(1,len(findings)),[label for ok,label in findings if not ok],earned,len(findings)


def _criterion(identifier:str,items:Iterable[tuple[bool,str]])->dict[str,Any]:
    score,failures,earned,total=_ratio(items)
    return {"criterion_id":identifier,"score":round(score,6),"evidence":"全部任务专属原子检查通过" if not failures else "未通过："+"；".join(failures[:24]),"method":"deterministic","raw":{"earned_units":earned,"total_units":total,"failure_count":len(failures)}}


def _formula(value:Any,*tokens:str)->bool:
    if not isinstance(value,str) or not value.startswith("="):return False
    normalized=re.sub(r"[\s'$]","",value).upper()
    return all(re.sub(r"[\s'$]","",token).upper() in normalized for token in tokens)


def _core(artifact:Path)->tuple[bool,list[str],list[str]]:
    errors=[];hashes=[]
    if not artifact.is_dir():return False,["输出目录不存在"],hashes
    for name in REQUIRED:
        path=artifact/name
        if not path.is_file() or path.stat().st_size==0:errors.append(f"核心交付缺失或为空：{name}");continue
        try:
            hashes.append(_sha256(path))
            if not zipfile.is_zipfile(path):raise ValueError("不是有效 OOXML ZIP")
            if name==XLSX:
                book=load_workbook(path,read_only=True,data_only=False)
                if not book.sheetnames:raise ValueError("Excel 无工作表")
                book.close()
            else:
                document=Document(path)
                if not document.paragraphs and not document.tables:raise ValueError("Word 无内容")
        except Exception as exc:errors.append(f"核心交付无法解析：{name}: {type(exc).__name__}: {exc}")
    source=_source()
    if not source.is_file() or _sha256(source)!=SOURCE_HASH:errors.append("冻结 transcripts.pdf 缺失或哈希错误")
    else:
        try:
            info=subprocess.run(["pdfinfo",str(source)],stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,timeout=60,check=False)
            match=re.search(r"^Pages:\s+(\d+)",info.stdout,flags=re.MULTILINE)
            if info.returncode!=0 or not match or int(match.group(1))!=1:errors.append("冻结 PDF 实际页数不是1")
        except Exception as exc:errors.append(f"冻结 PDF 无法检查：{exc}")
    return not errors,errors,hashes


def _fatal(errors:list[str],hashes:list[str])->dict[str,Any]:
    results=[_criterion(key,[(False,error) for error in errors]) for key in WEIGHTS]
    results.extend(_criterion(key,[(True,"未触发")]) for key in PENALTIES)
    return {"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":hashes or ["0"*64],"criterion_results":results,"reward":0.0,"passed":False,"hurdle_failures":["R001"],"penalties":[]}


def _recalculate(path:Path)->tuple[dict[str,list[list[Any]]]|None,str]:
    with tempfile.TemporaryDirectory(prefix="transcript-v2-recalc-") as raw:
        root=Path(raw);source=root/"source";output=root/"output";profile=root/"profile";source.mkdir();output.mkdir();copied=source/path.name;shutil.copy2(path,copied)
        env=os.environ.copy();env["HOME"]=str(root/"home");env["NO_PROXY"]="localhost,127.0.0.1,::1"
        try:completed=subprocess.run(["soffice","--headless",f"-env:UserInstallation=file://{profile}","--convert-to","xlsx","--outdir",str(output),str(copied)],stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,timeout=300,check=False,env=env)
        except Exception as exc:return None,str(exc)
        converted=output/path.name
        if completed.returncode!=0 or not converted.is_file():return None,f"退出={completed.returncode}；日志={completed.stdout[-400:]}"
        try:
            book=load_workbook(converted,data_only=True);values={sheet.title:[[cell.value for cell in row] for row in sheet.iter_rows()] for sheet in book.worksheets};book.close();return values,completed.stdout[-400:]
        except Exception as exc:return None,str(exc)


def _value(values:dict[str,list[list[Any]]]|None,sheet:str,row:int,column:int)->Any:
    try:return values[sheet][row-1][column-1] if values is not None else None
    except (KeyError,IndexError):return None


def _course_truth():
    result=[]
    for sequence,(year,term,course,tag,grade,credits) in enumerate(COURSES,1):
        points=V1.grade_points(grade);result.append((sequence,year,term,course,tag,grade,credits,points,points*credits))
    return result


def _term_truth():
    result=[];cum_credits=0.0;cum_points=0.0
    for year,term,displayed_term,displayed_cum,_ in TERMS:
        selected=[row for row in COURSES if row[0]==year and row[1]==term];credits=sum(row[5] for row in selected);points=sum(V1.grade_points(row[4])*row[5] for row in selected);cum_credits+=credits;cum_points+=points
        result.append({"year":year,"term":term,"count":len(selected),"credits":credits,"points":points,"term_gpa":points/credits,"displayed_term":displayed_term,"cum_credits":cum_credits,"cum_points":cum_points,"cum_gpa":cum_points/cum_credits,"displayed_cum":displayed_cum,"delta":displayed_cum-cum_points/cum_credits})
    return result


TERM_TRUTH=_term_truth()


def _workbook(artifact:Path)->tuple[dict[str,list[tuple[bool,str]]],dict[str,Any],list[str]]:
    findings={key:[] for key in ("R002","R003","R004","R005","R006","R007","R008")};context={};danger=[];path=artifact/XLSX
    try:
        package=V1.inspect_package(path);danger=[*package.get("macros",[]),*package.get("external",[]),*package.get("missing_targets",[])]
        book=load_workbook(path,data_only=False);values,recalc_log=_recalculate(path)
    except Exception as exc:
        for key in findings:findings[key].append((False,f"Excel 检查失败：{exc}"))
        return findings,context,danger
    expected_sheets=["Source_Profile","Course_Ledger","Term_Audit","Summary"]
    for index,name in enumerate(expected_sheets):findings["R008"].append((index<len(book.sheetnames) and book.sheetnames[index]==name,f"工作表第{index+1}项错误"))

    profile_rows=[
        ("Student","Teddy Roosevelt","Header — Student","Academic identifier retained; direct contact details excluded"),
        ("Issue_Date","October 25, 2018","Header — Date Issued","Source-reported date"),
        ("Credits_Earned",23.0,"Header — Credits Earned","Matches visible course-credit total"),
        ("Weighted_GPA",3.78,"Header — Weighted GPA","Source-reported; weighting policy not supplied"),
        ("Unweighted_GPA",3.65,"Header — Unweighted GPA","Matches independent A/B/C recalculation after rounding"),
        ("PDF_File_Type","Image-only","PDF object inspection","No extractable text layer"),
        ("Delivered_PDF_Pages",1,"PDF page tree","Only one page is present in the delivered file"),
        ("Printed_Page_Marker","1 of 2","Visible page footer","Printed marker claims a two-page source"),
        ("Printed_Total_Pages",2,"Visible page footer","Do not infer missing-page content"),
        ("Source_Completeness_Risk","Potentially incomplete","Delivered pages vs printed marker","Escalate before a completeness-sensitive decision"),
    ]
    profile=book["Source_Profile"] if "Source_Profile" in book.sheetnames else None
    for column,expected in enumerate(("Field","Extracted_Value","Evidence_Location","Audit_Note"),1):findings["R002"].append((profile is not None and _same(profile.cell(1,column).value,expected),f"Source_Profile 表头第{column}列错误"))
    for row,expected_row in enumerate(profile_rows,2):
        for column,expected in enumerate(expected_row,1):findings["R002"].append((_same(profile.cell(row,column).value if profile else None,expected),f"Source_Profile 第{row}行第{column}列错误"))
    context["profile"]={row[0]:profile.cell(i,2).value if profile else None for i,row in enumerate(profile_rows,2)}

    ledger=book["Course_Ledger"] if "Course_Ledger" in book.sheetnames else None
    for column,expected in enumerate(V1.LEDGER_HEADERS,1):findings["R003"].append((ledger is not None and _same(ledger.cell(1,column).value,expected),f"Course_Ledger 表头第{column}列错误"))
    ledger_context=[]
    for row,truth in enumerate(_course_truth(),2):
        for column in (1,2,3,4):findings["R003"].append((_same(ledger.cell(row,column).value if ledger else None,truth[column-1]),f"Course_Ledger 第{row}行第{column}列错误"))
        for column in (5,6,7):findings["R004"].append((_same(ledger.cell(row,column).value if ledger else None,truth[column-1]),f"Course_Ledger 第{row}行第{column}列错误"))
        formula_h=ledger.cell(row,8).value if ledger else None;formula_i=ledger.cell(row,9).value if ledger else None
        findings["R005"].extend([
            (_formula(formula_h,f"F{row}",'"A",4','"B",3','"C",2','"D",1','"F",0'),f"Course_Ledger!H{row} 成绩点公式错误"),
            (_formula(formula_i,f"G{row}*H{row}"),f"Course_Ledger!I{row} 质量分公式错误"),
            (_same(_value(values,"Course_Ledger",row,8),truth[7]),f"Course_Ledger!H{row} 重算错误：{recalc_log}"),
            (_same(_value(values,"Course_Ledger",row,9),truth[8]),f"Course_Ledger!I{row} 重算错误"),
        ])
        ledger_context.append([ledger.cell(row,column).value if ledger else None for column in range(1,8)]+[_value(values,"Course_Ledger",row,8),_value(values,"Course_Ledger",row,9)])
    context["ledger"]=ledger_context

    term=book["Term_Audit"] if "Term_Audit" in book.sheetnames else None
    for column,expected in enumerate(V1.TERM_LEFT_HEADERS,1):findings["R006"].append((term is not None and _same(term.cell(1,column).value,expected),f"Term_Audit 明细表头第{column}列错误"))
    for column,expected in enumerate(V1.TERM_RIGHT_HEADERS,1):findings["R006"].append((term is not None and _same(term.cell(18,column).value,expected),f"Term_Audit 累计表头第{column}列错误"))
    term_context=[]
    for offset,truth in enumerate(TERM_TRUTH):
        left=offset+2;right=offset+19
        expected_left=(truth["year"],truth["term"],truth["count"],truth["credits"],truth["points"],truth["term_gpa"],truth["displayed_term"])
        expected_right=(truth["year"],truth["term"],truth["cum_credits"],truth["cum_points"],truth["cum_gpa"],truth["displayed_cum"],truth["delta"])
        for column,target in enumerate(expected_left,1):
            actual=_value(values,"Term_Audit",left,column) if column in (3,4,5,6) else (term.cell(left,column).value if term else None)
            findings["R006"].append((_same(actual,target),f"Term_Audit 第{left}行第{column}列错误"))
        for column,target in enumerate(expected_right,1):
            actual=_value(values,"Term_Audit",right,column) if column in (3,4,5,7) else (term.cell(right,column).value if term else None)
            findings["R006"].append((_same(actual,target),f"Term_Audit 第{right}行第{column}列错误"))
        for column,tokens in ((3,("COUNTIFS","COURSE_LEDGER")),(4,("SUMIFS","COURSE_LEDGER")),(5,("SUMIFS","COURSE_LEDGER")),(6,("IFERROR",f"E{left}/D{left}"))):findings["R006"].append((_formula(term.cell(left,column).value if term else None,*tokens),f"Term_Audit!{chr(64+column)}{left} 公式错误"))
        for column,tokens in ((3,("SUM",f"$D$2:D{left}")),(4,("SUM",f"$E$2:E{left}")),(5,("IFERROR",f"D{right}/C{right}")),(7,(f"F{right}-E{right}",))):findings["R006"].append((_formula(term.cell(right,column).value if term else None,*tokens),f"Term_Audit!{chr(64+column)}{right} 公式错误"))
        term_context.append((truth["year"],truth["term"],truth["term_gpa"],truth["displayed_term"],truth["cum_gpa"],truth["displayed_cum"]))
    context["terms"]=term_context

    summary=book["Summary"] if "Summary" in book.sheetnames else None
    summary_values=[23,23.0,17,4,2,84.0,84/23,3.65,3.65,0.0,3.78,3.78,0.0,1,2,"Potentially incomplete"]
    summary_tokens=[("COUNTA","COURSE_LEDGER"),("SUM","COURSE_LEDGER"),("COUNTIF",'"A"'),("COUNTIF",'"B"'),("COUNTIF",'"C"'),("SUM","COURSE_LEDGER"),("B7/B3",),("ROUND","B8"),("SOURCE_PROFILE","B6"),("B9-B10",),("SOURCE_PROFILE","B5"),("TERM_AUDIT","F33"),("B13-B12",),("SOURCE_PROFILE","B8"),("SOURCE_PROFILE","B10"),("IF","B15<B16","POTENTIALLYINCOMPLETE")]
    summary_context={}
    for row,(label,target,tokens) in enumerate(zip(V1.SUMMARY_LABELS,summary_values,summary_tokens,strict=True),2):
        findings["R007"].extend([
            (_same(summary.cell(row,1).value if summary else None,label),f"Summary!A{row} 标签错误"),
            (_formula(summary.cell(row,2).value if summary else None,*tokens),f"Summary!B{row} 公式错误"),
            (_same(_value(values,"Summary",row,2),target),f"Summary!B{row} 重算错误"),
        ]);summary_context[label]=_value(values,"Summary",row,2)
    context["summary"]=summary_context

    expected_tables={"Source_Profile":["TranscriptSourceProfileTable"],"Course_Ledger":["TranscriptCourseLedgerTable"],"Term_Audit":["TranscriptTermDetailTable","TranscriptCumulativeAuditTable"],"Summary":["TranscriptAuditSummaryTable"]}
    for sheet_name,names in expected_tables.items():
        sheet=book[sheet_name] if sheet_name in book.sheetnames else None
        findings["R008"].extend([(sheet is not None and set(sheet.tables)==set(names),f"{sheet_name} 原生表错误"),(sheet is not None and sheet.freeze_panes=="A2",f"{sheet_name} 未冻结首行"),(sheet is not None and sheet.auto_filter.ref is not None,f"{sheet_name} 缺筛选")])
    findings["R008"].extend([(summary is not None and len(summary._charts)==1,"Summary 缺可编辑折线图"),(sum(len(sheet.conditional_formatting) for sheet in book.worksheets)>=5,"条件格式不足"),(not any(len(sheet._images) for sheet in book.worksheets),"Excel 使用图片替代可编辑内容"),(not danger,f"OOXML 含危险/断裂对象：{danger}")])
    book.close();return findings,context,danger


def _table_rows(table)->list[list[str]]:return [[cell.text for cell in row.cells] for row in table.rows]


def _word(artifact:Path)->tuple[dict[str,list[tuple[bool,str]]],dict[str,Any],list[str],list[str]]:
    findings={key:[] for key in ("R009","R010","R011")};path=artifact/DOCX;document=Document(path);text=V1.word_text(document);lower=_norm(text);unsupported=V1.unsupported_claims(text);privacy=V1.artifact_privacy_findings(artifact)
    headings=[V1.normalize(paragraph.text) for paragraph in document.paragraphs if str(paragraph.style.name).startswith("Heading")]
    findings["R009"].append(("Transcript GPA Audit Report" in text,"Word 标题错误"))
    for heading in ("Source Completeness","Reported Academic Profile","Course Extraction and Recalculation","Term-Level Reconciliation","Weighted GPA Evidence Boundary","Audit Conclusion"):findings["R009"].append((heading in headings,f"Word 缺章节 {heading}"))
    for token in ("image-only PDF","1 actual page","1 of 2","potentially incomplete","No content from the unprovided page is inferred"):findings["R009"].append((token.casefold() in lower,f"Word 完整性说明缺 {token}"))
    findings["R009"].extend([(len(document.tables)==4,f"Word 原生表数错误：{len(document.tables)}"),(not document.inline_shapes,"Word 用图片替代表格")])

    fact_tokens=("Teddy Roosevelt","October 25, 2018","23 courses","15 terms","23.00 credits","17 A","4 B","2 C","84.00 unweighted quality points","84/23 = 3.652173913","reported unweighted GPA","3.65","reported weighted GPA","3.78","2017/2018 Spring","Drama (DE)","World Literature (AP)","3.50","5.00","difference of 1.50")
    for token in fact_tokens:findings["R010"].append((token.casefold() in lower,f"Word 缺事实 {token}"))
    findings["R010"].extend([
        ("the visible page contains 23 courses across 15 terms and 23.00 credits." in lower,"Word 课程/学期/学分主句错误"),
        ("its unweighted recalculation is 3.50, while the pdf displays a term gpa of 5.00, a difference of 1.50." in lower,"Word 末学期对账主句错误"),
    ])
    expected_profile_tables=[
        [["Courses","Credits","Reported unweighted","Reported weighted"],["23","23.00","3.65 / 4.00","3.78 / 4.00"]],
        [["Field","Source-reported value","Audit treatment"],["Student","Teddy Roosevelt","Retained as the academic identifier"],["Credits earned","23.00","Reconciled to the visible course ledger"],["Unweighted GPA","3.65 / 4.00","Independently recalculated"],["Weighted GPA","3.78 / 4.00","Retained as reported; policy unavailable"]],
    ]
    for table_index,expected_rows in enumerate(expected_profile_tables):
        profile_table=document.tables[table_index] if table_index<len(document.tables) else None
        for row_index,expected_row in enumerate(expected_rows):
            for column_index,target in enumerate(expected_row):
                actual=profile_table.rows[row_index].cells[column_index].text if profile_table is not None and row_index<len(profile_table.rows) else None
                findings["R010"].append((_same(actual,target),f"Word 概要表{table_index+1}第{row_index+1}行第{column_index+1}列错误"))
    expected_term_rows=[["Period","Courses","Credits","Calc term","Displayed term","Calc cumulative","Displayed cumulative"]]
    for truth in TERM_TRUTH:expected_term_rows.append([f"{truth['year']} {truth['term']}",str(truth['count']),f"{truth['credits']:.2f}",f"{truth['term_gpa']:.2f}",f"{truth['displayed_term']:.2f}",f"{truth['cum_gpa']:.2f}",f"{truth['displayed_cum']:.2f}"])
    table=document.tables[2] if len(document.tables)>2 else None
    findings["R010"].append((table is not None and len(table.rows)==16,"Word 学期表行数错误"))
    for row_index,expected_row in enumerate(expected_term_rows):
        for column_index,target in enumerate(expected_row):
            actual=table.rows[row_index].cells[column_index].text if table is not None and row_index<len(table.rows) and column_index<len(table.rows[row_index].cells) else None
            findings["R010"].append((_same(actual,target),f"Word 学期表第{row_index+1}行第{column_index+1}列错误"))
    expected_summary_table=[
        ["Check","Result","Boundary"],
        ["Unweighted course-level GPA","84.00 / 23.00 = 3.652173913 → 3.65","Recomputed and reconciled"],
        ["Reported weighted GPA","3.78","Retained; not independently reconstructed"],
        ["Final displayed cumulative GPA","3.78","Matches the reported weighted value"],
        ["Source completeness","1 delivered page vs printed 1 of 2","Potentially incomplete"],
    ]
    boundary_table=document.tables[3] if len(document.tables)>3 else None
    for row_index,expected_row in enumerate(expected_summary_table):
        for column_index,target in enumerate(expected_row):
            actual=boundary_table.rows[row_index].cells[column_index].text if boundary_table is not None and row_index<len(boundary_table.rows) else None
            findings["R011"].append((_same(actual,target),f"Word 边界表第{row_index+1}行第{column_index+1}列错误"))
    for token in ("weighting policy","unprovided second page","cannot independently reconstruct","does not assume conventional DE/AP bonuses","reported-but-unverified","not an authenticity, graduation or eligibility determination"):findings["R011"].append((token.casefold() in lower,f"Word 证据边界缺 {token}"))
    findings["R011"].append((not unsupported,f"Word 含无来源断言：{unsupported}"));findings["R011"].append((not privacy,f"交付物泄漏隐私：{privacy}"))
    return findings,{"text":text,"tables":[_table_rows(table) for table in document.tables]},unsupported,privacy


def _render(artifact:Path)->dict[str,Any]:return V1.render_office(artifact)


def evaluate(artifact:Path,scoring_version:str|None=None)->dict[str,Any]:
    if (scoring_version or os.environ.get("SCORING_VERSION","v2")).casefold()=="v1":return V1.evaluate(artifact)
    core_ok,errors,hashes=_core(artifact)
    if not core_ok:return _fatal(errors,hashes)
    criteria={key:[] for key in WEIGHTS};criteria["R001"].append((True,"两份核心交付与冻结输入均可解析"))
    workbook,workbook_context,danger=_workbook(artifact)
    for key,items in workbook.items():criteria[key].extend(items)
    word,word_context,unsupported,privacy=_word(artifact)
    for key,items in word.items():criteria[key].extend(items)
    summary=workbook_context.get("summary",{});tables=word_context.get("tables",[])
    expected_cross={"Extracted_Courses":23,"Extracted_Credits":23,"Grade_A_Count":17,"Grade_B_Count":4,"Grade_C_Count":2,"Unweighted_Quality_Points":84,"Calculated_Unweighted_GPA":84/23,"Reported_Unweighted_GPA":3.65,"Reported_Weighted_GPA":3.78,"Final_Displayed_Cumulative_GPA":3.78,"Delivered_PDF_Pages":1,"Printed_Total_Pages":2,"Source_Completeness_Flag":"Potentially incomplete"}
    for label,target in expected_cross.items():criteria["R012"].append((_same(summary.get(label),target),f"Excel 汇总 {label} 错误"))
    text=word_context.get("text","")
    for token in ("Teddy Roosevelt","23.00 credits","17 A","4 B","2 C","84.00","3.652173913","3.65","3.78","1 actual page","1 of 2","potentially incomplete"):criteria["R012"].append((token.casefold() in text.casefold(),f"Excel/Word 一致性缺 {token}"))
    render=_render(artifact);criteria["R013"].extend([
        (not render.get("errors"),f"LibreOffice 渲染错误：{render.get('errors')}"),
        (render.get("xlsx_pages")==4,f"Excel 页数错误：{render.get('xlsx_pages')}"),
        (render.get("docx_pages")==3,f"Word 页数错误：{render.get('docx_pages')}"),
    ])
    extra=sorted(path.name for path in artifact.iterdir() if path.name not in REQUIRED);p1_bad=len(unsupported)+len(privacy);p1_score=min(1.0,p1_bad/2.0);p2_bad=bool(extra or danger)
    positive=[_criterion(key,criteria[key]) for key in WEIGHTS];base=sum(WEIGHTS[item["criterion_id"]]*float(item["score"]) for item in positive)/100.0
    applied=[]
    if p1_score:applied.append({"id":"P001","amount":round(.12*p1_score,6),"reason":"无来源结论或隐私泄漏"})
    if p2_bad:applied.append({"id":"P002","amount":.04,"reason":f"额外交付或危险对象：extra={extra}"})
    reward=round(max(0.0,base-sum(item["amount"] for item in applied)),6)
    positive.extend([{"criterion_id":"P001","score":round(p1_score,6),"evidence":"未触发" if not p1_score else "检测到无来源断言/隐私","method":"deterministic","raw":{"unsupported":unsupported,"privacy":privacy}},{"criterion_id":"P002","score":float(p2_bad),"evidence":"未触发" if not p2_bad else "检测到额外交付/危险对象","method":"deterministic","raw":{"extra":extra,"danger":danger}}])
    return {"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":hashes,"criterion_results":positive,"base_score":round(base,6),"reward":reward,"passed":reward>=.8,"hurdle_failures":[],"penalties":applied}


def main()->None:
    parser=argparse.ArgumentParser();parser.add_argument("--artifact-dir",type=Path,required=True);parser.add_argument("--result",type=Path,required=True);parser.add_argument("--reward-file",type=Path,required=True);parser.add_argument("--scoring-version",choices=("v1","v2"),default=os.environ.get("SCORING_VERSION","v2"));args=parser.parse_args()
    try:result=evaluate(args.artifact_dir,args.scoring_version)
    except Exception as exc:result={"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"reward":0.0,"passed":False,"fatal_error":f"{type(exc).__name__}: {exc}","hurdle_failures":["R001"]}
    args.result.parent.mkdir(parents=True,exist_ok=True);args.reward_file.parent.mkdir(parents=True,exist_ok=True);args.result.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n",encoding="utf-8");args.reward_file.write_text(f"{float(result['reward']):.6f}\n",encoding="utf-8");print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=="__main__":main()
