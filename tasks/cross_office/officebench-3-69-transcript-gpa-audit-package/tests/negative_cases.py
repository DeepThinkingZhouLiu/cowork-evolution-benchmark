#!/usr/bin/env python3
"""成绩单 GPA 审计包的 v2 任务专属分层负例。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Callable

from docx import Document
from openpyxl import Workbook, load_workbook


ROOT=Path(__file__).resolve().parents[1];GOLDEN=ROOT/"solution"/"Golden";INPUTS=ROOT/"environment"/"input_files"
REQUIRED=("Transcript_GPA_Audit.xlsx","Transcript_GPA_Audit_Report.docx")
Mutator=Callable[[Path],None]


def _load_judge():
    path=Path(__file__).with_name("judge.py");spec=importlib.util.spec_from_file_location("transcript_v2_judge_for_negatives",path)
    if spec is None or spec.loader is None:raise RuntimeError("无法载入 v2 Judge")
    module=importlib.util.module_from_spec(spec);sys.modules[spec.name]=module;spec.loader.exec_module(module);return module


JUDGE=_load_judge()


def _copy(target:Path)->None:
    target.mkdir(parents=True,exist_ok=True)
    for name in REQUIRED:shutil.copy2(GOLDEN/name,target/name)


def _excel(sheet:str,cell:str,value:object)->Mutator:
    def apply(root:Path)->None:
        path=root/REQUIRED[0];book=load_workbook(path);book[sheet][cell]=value;book.save(path);book.close()
    return apply


def _remove_chart(root:Path)->None:
    path=root/REQUIRED[0];book=load_workbook(path);book["Summary"]._charts=[];book.save(path);book.close()


def _remove_freeze(root:Path)->None:
    path=root/REQUIRED[0];book=load_workbook(path);book["Course_Ledger"].freeze_panes=None;book.save(path);book.close()


def _remove_table(root:Path)->None:
    path=root/REQUIRED[0];book=load_workbook(path);del book["Course_Ledger"].tables["TranscriptCourseLedgerTable"];book.save(path);book.close()


def _word_replace(old:str,new:str)->Mutator:
    def apply(root:Path)->None:
        path=root/REQUIRED[1];document=Document(path);locations=list(document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:locations.extend(cell.paragraphs)
        for paragraph in locations:
            if old in paragraph.text:
                paragraph.text=paragraph.text.replace(old,new,1);document.save(path);return
        raise AssertionError(f"Word 未找到：{old}")
    return apply


def _word_table(table_index:int,row_index:int,column_index:int,value:str)->Mutator:
    def apply(root:Path)->None:
        path=root/REQUIRED[1];document=Document(path);document.tables[table_index].rows[row_index].cells[column_index].text=value;document.save(path)
    return apply


def _word_remove_table(root:Path)->None:
    path=root/REQUIRED[1];document=Document(path);table=document.tables[2];table._element.getparent().remove(table._element);document.save(path)


def _word_add(text:str)->Mutator:
    def apply(root:Path)->None:
        path=root/REQUIRED[1];document=Document(path);document.add_paragraph(text);document.save(path)
    return apply


def _moderate_reconciliation_domain(root:Path)->None:
    path=root/REQUIRED[0];book=load_workbook(path)
    term=book["Term_Audit"]
    for row in range(2,17):
        for column in (3,4,5,6):term.cell(row,column).value=0
        term.cell(row,7).value=3.0
    for row in range(19,34):
        for column in (3,4,5,7):term.cell(row,column).value=0
        term.cell(row,6).value=3.0
    summary=book["Summary"]
    for row in range(2,17):summary.cell(row,2).value=0
    summary["B17"]="Complete";book.save(path);book.close()
    doc=root/REQUIRED[1];document=Document(doc);table=document.tables[2]
    for row in range(1,16):
        for column in range(1,7):table.rows[row].cells[column].text="0.00"
    boundary=document.tables[3]
    for row in range(1,5):boundary.rows[row].cells[1].text="0.00";boundary.rows[row].cells[2].text="Verified"
    document.paragraphs[6].text="The transcript has 0 courses, 0 credits and a GPA of 0.00."
    document.paragraphs[9].text="All term values are 0.00."
    document.paragraphs[12].text="The source weighted GPA is treated as 3.00 and independently verified under the school policy."
    document.save(doc)


def _severe_valid_incomplete(root:Path)->None:
    book=Workbook();book.remove(book.active)
    headers={"Source_Profile":["Field","Extracted_Value","Evidence_Location","Audit_Note"],"Course_Ledger":JUDGE.V1.LEDGER_HEADERS,"Term_Audit":JUDGE.V1.TERM_LEFT_HEADERS,"Summary":["Metric","Value","Interpretation"]}
    for name,row in headers.items():sheet=book.create_sheet(name);sheet.append(row)
    book.save(root/REQUIRED[0]);book.close()
    document=Document();document.add_paragraph("Transcript GPA Audit Report",style="Title");document.add_paragraph("Evidence review is incomplete.");document.save(root/REQUIRED[1])


def _input_only(root:Path)->None:
    for path in list(root.iterdir()):
        if path.is_file():path.unlink()
    for source in INPUTS.iterdir():
        if source.is_file():shutil.copy2(source,root/source.name)


PARTIALS:dict[str,tuple[Mutator,bool]]={
    "profile_student_wrong":(_excel("Source_Profile","B2","Theodore Roosevelt"),False),
    "profile_issue_date_wrong":(_excel("Source_Profile","B3","October 26, 2018"),False),
    "profile_credits_wrong":(_excel("Source_Profile","B4",24),False),
    "profile_weighted_gpa_wrong":(_excel("Source_Profile","B5",3.9),False),
    "profile_unweighted_gpa_wrong":(_excel("Source_Profile","B6",3.7),False),
    "profile_delivered_pages_wrong":(_excel("Source_Profile","B8",2),False),
    "profile_printed_marker_wrong":(_excel("Source_Profile","B9","1 of 1"),False),
    "profile_completeness_wrong":(_excel("Source_Profile","B11","Complete"),False),
    "course_sequence_wrong":(_excel("Course_Ledger","A2",2),False),
    "course_year_wrong":(_excel("Course_Ledger","B2","2015/2016"),False),
    "course_term_wrong":(_excel("Course_Ledger","C2","Winter"),False),
    "course_name_wrong":(_excel("Course_Ledger","D24","British Literature"),False),
    "course_tag_missing":(_excel("Course_Ledger","E23",None),False),
    "course_tag_invented":(_excel("Course_Ledger","E2","AP"),False),
    "course_grade_wrong":(_excel("Course_Ledger","F15","B"),False),
    "course_credit_wrong":(_excel("Course_Ledger","G2",2),False),
    "grade_points_flattened":(_excel("Course_Ledger","H2",4),False),
    "grade_points_scale_wrong":(_excel("Course_Ledger","H24",'=IF(F24="B",4,0)'),False),
    "quality_formula_flattened":(_excel("Course_Ledger","I2",4),False),
    "quality_formula_wrong_reference":(_excel("Course_Ledger","I2","=G2*H3"),False),
    "term_displayed_gpa_wrong":(_excel("Term_Audit","G16",3.5),False),
    "term_count_formula_flattened":(_excel("Term_Audit","C2",2),False),
    "term_credit_formula_wrong":(_excel("Term_Audit","D2","=SUM(Course_Ledger!G2:G24)"),False),
    "term_gpa_formula_flattened":(_excel("Term_Audit","F2",4),False),
    "cumulative_credits_formula_flattened":(_excel("Term_Audit","C19",2),False),
    "cumulative_gpa_formula_wrong":(_excel("Term_Audit","E33","=4"),False),
    "cumulative_delta_formula_flattened":(_excel("Term_Audit","G33",0.127826086956522),False),
    "summary_course_count_flattened":(_excel("Summary","B2",23),False),
    "summary_gpa_wrong_formula":(_excel("Summary","B8","=4"),False),
    "summary_completeness_flattened":(_excel("Summary","B17","Potentially incomplete"),False),
    "native_chart_removed":(_remove_chart,False),
    "native_freeze_removed":(_remove_freeze,False),
    "native_table_removed":(_remove_table,False),
    "word_heading_wrong":(_word_replace("Source Completeness","Source Review"),False),
    "word_course_count_wrong":(_word_replace("23 courses","24 courses"),False),
    "word_credit_conflict":(_word_replace("23.00 credits","24.00 credits"),False),
    "word_last_term_calc_wrong":(_word_table(2,15,3,"4.00"),False),
    "word_last_term_display_wrong":(_word_table(2,15,4,"3.50"),False),
    "word_boundary_table_wrong":(_word_table(3,2,1,"4.00"),False),
    "word_source_complete_claim":(_word_add("The delivered source is complete; no page is missing."),False),
    "word_invented_weight_rule":(_word_add("The school's weighting policy adds 1.0 point to every AP and DE grade."),False),
    "word_authenticity_claim":(_word_add("This transcript is authentic and officially valid."),False),
    "word_privacy_address":(_word_add("Student address: 1900 Pennsylvania Ave"),False),
    "word_privacy_phone":(_word_add("Student phone: (202) 456-1111"),False),
    "word_term_table_removed":(_word_remove_table,False),
    "unexpected_debug_file":(lambda root:(root/"debug.txt").write_text("temporary\n",encoding="utf-8"),False),
    "moderate_reconciliation_domain_conflict":(_moderate_reconciliation_domain,True),
    "severe_valid_but_incomplete_package":(_severe_valid_incomplete,True),
}


def main()->None:
    parser=argparse.ArgumentParser();parser.add_argument("--report",type=Path);args=parser.parse_args();golden_recalc,_=JUDGE._recalculate(GOLDEN/REQUIRED[0]);golden_render=JUDGE._render(GOLDEN);original_recalc,original_render=JUDGE._recalculate,JUDGE._render
    assert JUDGE.evaluate(GOLDEN)["reward"]==1.0
    def fast_recalc(path:Path):
        book=load_workbook(path,data_only=False);values={}
        for sheet in book.worksheets:
            rows=[]
            for row_index,row in enumerate(sheet.iter_rows(),1):
                current=[]
                for column_index,cell in enumerate(row,1):
                    value=cell.value
                    if isinstance(value,str) and value.startswith("="):value=JUDGE._value(golden_recalc,sheet.title,row_index,column_index)
                    current.append(value)
                rows.append(current)
            values[sheet.title]=rows
        book.close();return values,"task-specific cached Golden formula values"
    def evaluate_case(case:Path,real:bool)->float:
        if real:return float(JUDGE.evaluate(case)["reward"])
        JUDGE._recalculate=fast_recalc;JUDGE._render=lambda _artifact:golden_render
        try:return float(JUDGE.evaluate(case)["reward"])
        finally:JUDGE._recalculate,JUDGE._render=original_recalc,original_render
    records=[]
    with tempfile.TemporaryDirectory(prefix="transcript-v2-negatives-") as raw:
        base=Path(raw)
        for name in REQUIRED:
            for mode in ("missing","empty","corrupt"):
                case=base/f"fatal-{mode}-{name}";_copy(case);target=case/name
                if mode=="missing":target.unlink()
                elif mode=="empty":target.write_bytes(b"")
                else:target.write_bytes(b"not-a-valid-task-artifact")
                reward=evaluate_case(case,False);assert reward==0,(mode,name,reward);records.append({"case":f"{mode}:{name}","layer":"fatal_zero","reward":reward})
        empty=base/"fatal-empty";empty.mkdir();reward=evaluate_case(empty,False);assert reward==0;records.append({"case":"empty-output","layer":"fatal_zero","reward":reward})
        input_only=base/"fatal-input-only";_copy(input_only);_input_only(input_only);reward=evaluate_case(input_only,False);assert reward==0;records.append({"case":"unmodified-input-only","layer":"fatal_zero","reward":reward})
        for name,(mutator,real) in PARTIALS.items():
            case=base/f"partial-{name}";_copy(case);mutator(case);reward=evaluate_case(case,real);assert 0<reward<1,(name,reward);records.append({"case":name,"layer":"partial","reward":reward})
    partial=[float(record["reward"]) for record in records if record["layer"]=="partial"];by_name={str(record["case"]):float(record["reward"]) for record in records}
    assert by_name["moderate_reconciliation_domain_conflict"]<.8
    assert 0<by_name["severe_valid_but_incomplete_package"]<by_name["moderate_reconciliation_domain_conflict"]
    report={"total":len(records),"fatal_zero":sum(record["layer"]=="fatal_zero" for record in records),"partial_continuous":len(partial),"partial_reward_range":[min(partial),max(partial)],"moderate_reconciliation_domain_conflict":by_name["moderate_reconciliation_domain_conflict"],"severe_valid_but_incomplete":by_name["severe_valid_but_incomplete_package"],"cases":records}
    if args.report:args.report.parent.mkdir(parents=True,exist_ok=True);args.report.write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=="__main__":main()
