#!/usr/bin/env python3
"""构造并回放 OfficeBench 3-69 的任务专属负例。"""

from __future__ import annotations

import argparse
import shutil
import tempfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from judge import DOCX_NAME, REQUIRED_FILES, XLSX_NAME, evaluate


CASES = (
    "missing-docx",
    "formulas-flattened",
    "wrong-course-name",
    "wrong-grade",
    "wrong-program-tag",
    "missing-course",
    "wrong-term-displayed-gpa",
    "wrong-summary-gpa",
    "source-completeness-denied",
    "invented-weight-rule",
    "word-cross-file-credit-conflict",
    "privacy-leak",
    "corrupt-docx",
    "extra-debug-file",
    "native-chart-removed",
)


def copy_golden(golden: Path, case_dir: Path) -> None:
    case_dir.mkdir(parents=True, exist_ok=True)
    for name in REQUIRED_FILES:
        shutil.copy2(golden / name, case_dir / name)


def replace_text(path: Path, old: str, new: str) -> None:
    document = Document(path)
    locations = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                locations.extend(cell.paragraphs)
    replaced = False
    for paragraph in locations:
        if old not in paragraph.text:
            continue
        replaced = True
        if any(old in run.text for run in paragraph.runs):
            for run in paragraph.runs:
                run.text = run.text.replace(old, new)
        else:
            paragraph.text = paragraph.text.replace(old, new)
    if not replaced:
        raise ValueError(f"未找到待替换文本：{old}")
    document.save(path)


def mutate(case: str, case_dir: Path) -> None:
    workbook_path = case_dir / XLSX_NAME
    document_path = case_dir / DOCX_NAME
    if case == "missing-docx":
        document_path.unlink()
    elif case == "formulas-flattened":
        formulas = load_workbook(workbook_path, data_only=False)
        values = load_workbook(workbook_path, data_only=True)
        for formula_sheet, value_sheet in zip(formulas.worksheets, values.worksheets, strict=True):
            for row in formula_sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        cell.value = value_sheet[cell.coordinate].value
        formulas.save(workbook_path)
    elif case == "wrong-course-name":
        workbook = load_workbook(workbook_path)
        workbook["Course_Ledger"]["D24"] = "British Literature"
        workbook.save(workbook_path)
    elif case == "wrong-grade":
        workbook = load_workbook(workbook_path)
        workbook["Course_Ledger"]["F15"] = "B"
        workbook.save(workbook_path)
    elif case == "wrong-program-tag":
        workbook = load_workbook(workbook_path)
        workbook["Course_Ledger"]["E23"] = "AP"
        workbook.save(workbook_path)
    elif case == "missing-course":
        workbook = load_workbook(workbook_path)
        workbook["Course_Ledger"].delete_rows(24)
        workbook.save(workbook_path)
    elif case == "wrong-term-displayed-gpa":
        workbook = load_workbook(workbook_path)
        workbook["Term_Audit"]["G16"] = 3.5
        workbook.save(workbook_path)
    elif case == "wrong-summary-gpa":
        workbook = load_workbook(workbook_path)
        workbook["Summary"]["B8"] = "=4"
        workbook.save(workbook_path)
    elif case == "source-completeness-denied":
        replace_text(document_path, "potentially incomplete", "complete")
        document = Document(document_path)
        document.add_paragraph("The delivered source is complete; no page is missing.")
        document.save(document_path)
    elif case == "invented-weight-rule":
        document = Document(document_path)
        document.add_paragraph("The school's weighting policy adds 1.0 point to every AP and DE grade.")
        document.save(document_path)
    elif case == "word-cross-file-credit-conflict":
        replace_text(document_path, "23.00 credits", "24.00 credits")
    elif case == "privacy-leak":
        document = Document(document_path)
        document.add_paragraph("Student address: 1900 Pennsylvania Ave")
        document.save(document_path)
    elif case == "corrupt-docx":
        document_path.write_bytes(b"not-a-docx")
    elif case == "extra-debug-file":
        (case_dir / "debug.txt").write_text("temporary", encoding="utf-8")
    elif case == "native-chart-removed":
        workbook = load_workbook(workbook_path)
        workbook["Summary"]._charts = []
        workbook.save(workbook_path)
    else:
        raise ValueError(case)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--case", choices=CASES, action="append", required=True)
    args = parser.parse_args()
    failures: list[dict[str, object]] = []
    for case in args.case:
        with tempfile.TemporaryDirectory(prefix=f"transcript-negative-{case}-") as tmp:
            case_dir = Path(tmp)
            copy_golden(args.golden, case_dir)
            mutate(case, case_dir)
            result = evaluate(case_dir)
            print(
                f"{case}\treward={result['reward']:.6f}\tpassed={result['passed']}"
                f"\thurdles={','.join(result['hurdle_failures'])}",
                flush=True,
            )
            if result["passed"] or float(result["reward"]) >= 0.8:
                failures.append({"case": case, "result": result})
    if failures:
        raise SystemExit(f"负例错误通过：{[item['case'] for item in failures]}")


if __name__ == "__main__":
    main()
