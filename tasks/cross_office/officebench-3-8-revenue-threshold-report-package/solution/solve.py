#!/usr/bin/env python3
"""根据 OfficeBench 3-8 的冻结收入表生成本题两份相互一致的 Golden。"""

from __future__ import annotations

import argparse
import hashlib
import shutil
import subprocess
import tempfile
import zipfile
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


XLSX_NAME = "Revenue_Filtered_Analysis.xlsx"
DOCX_NAME = "Revenue_Threshold_Report.docx"
THRESHOLD = 2_000_000
FIXED_DT = datetime(2026, 9, 1, 5, 45, 0, tzinfo=timezone.utc)

NAVY = "17324D"
BLUE = "2A6F97"
TEAL = "168AAD"
GOLD = "D9A441"
WHITE = "FFFFFF"
PALE_BLUE = "EAF3F8"
PALE_GREEN = "E7F4EC"
PALE_RED = "FBE9E7"
PALE_GOLD = "FFF5DB"
LINE = "CCD8E2"
GRAY = "607080"


def money(value: Decimal | int) -> str:
    decimal_value = Decimal(value)
    rounded = decimal_value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return f"-${abs(rounded):,.2f}" if rounded < 0 else f"${rounded:,.2f}"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_source(input_dir: Path) -> list[tuple[int, int]]:
    source = input_dir / "company_revenues.xlsx"
    if not source.is_file():
        raise FileNotFoundError(source)
    workbook = load_workbook(source, data_only=False, read_only=True)
    if workbook.sheetnames != ["Sheet1"]:
        raise ValueError(f"输入工作表错误：{workbook.sheetnames}")
    sheet = workbook["Sheet1"]
    headers = [str(cell.value or "").strip() for cell in sheet[1]]
    if headers != ["Year", "Revenue ($)"]:
        raise ValueError(f"输入表头错误：{headers}")
    rows = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in row):
            continue
        rows.append((int(row[0]), int(row[1])))
    workbook.close()
    if rows != sorted(rows) or len(rows) != 20 or [year for year, _ in rows] != list(range(2004, 2024)):
        raise ValueError("输入必须恰好包含 2004–2023 的 20 条升序记录")
    return rows


def style_sheet(sheet, widths: dict[str, float], orientation: str = "landscape") -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 32
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color=LINE)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=9, color=NAVY)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.page_setup.orientation = orientation
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:1"
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.35
    sheet.page_margins.bottom = 0.35


def add_native_table(sheet, name: str) -> None:
    table = Table(displayName=name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    sheet.add_table(table)


def build_workbook(rows: list[tuple[int, int]], output_path: Path) -> None:
    retained = [(year, revenue) for year, revenue in rows if revenue >= THRESHOLD]
    removed = [(year, revenue) for year, revenue in rows if revenue < THRESHOLD]
    if len(retained) != 17 or removed != [(2006, 1193225), (2009, 1791112), (2017, 1463500)]:
        raise ValueError("冻结真值与预期不一致")

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.properties.creator = "EvoCowork OfficeBench adapter"
    workbook.properties.title = "Revenue Threshold Analysis"
    workbook.properties.created = FIXED_DT.replace(tzinfo=None)
    workbook.properties.modified = FIXED_DT.replace(tzinfo=None)

    source = workbook.create_sheet("Source_Data")
    source.append(["Year", "Revenue_USD"])
    for year, revenue in rows:
        source.append([year, revenue])
    style_sheet(source, {"A": 14, "B": 22}, orientation="portrait")
    for row_index in range(2, 22):
        source.cell(row_index, 2).number_format = '$#,##0'
    add_native_table(source, "RevenueSourceTable")
    source.print_area = "A1:B21"

    audit = workbook.create_sheet("Threshold_Audit")
    audit.append(["Year", "Revenue_USD", "Threshold_USD", "Status", "Margin_to_Threshold", "Revenue_Rank"])
    for row_index, (year, revenue) in enumerate(rows, start=2):
        audit.append(
            [year, revenue, THRESHOLD, f'=IF(B{row_index}>=C{row_index},"Retain","Remove")', f"=B{row_index}-C{row_index}", f"=RANK(B{row_index},$B$2:$B$21,0)"]
        )
    style_sheet(audit, {"A": 12, "B": 18, "C": 18, "D": 13, "E": 22, "F": 14})
    for row_index in range(2, 22):
        for column in (2, 3, 5):
            audit.cell(row_index, column).number_format = '$#,##0'
    audit.conditional_formatting.add("D2:D21", FormulaRule(formula=['D2="Retain"'], fill=PatternFill("solid", fgColor=PALE_GREEN)))
    audit.conditional_formatting.add("D2:D21", FormulaRule(formula=['D2="Remove"'], fill=PatternFill("solid", fgColor=PALE_RED)))
    add_native_table(audit, "RevenueThresholdAuditTable")
    audit.print_area = "A1:F21"

    retained_sheet = workbook.create_sheet("Retained_Years")
    retained_sheet.append(["Year", "Revenue_USD", "Margin_to_Threshold", "Change_vs_Previous_Retained", "Change_Pct_vs_Previous_Retained"])
    for row_index, (year, revenue) in enumerate(retained, start=2):
        change = None if row_index == 2 else f"=B{row_index}-B{row_index-1}"
        change_pct = None if row_index == 2 else f'=IFERROR(B{row_index}/B{row_index-1}-1,"")'
        retained_sheet.append([year, revenue, f"=B{row_index}-{THRESHOLD}", change, change_pct])
    style_sheet(retained_sheet, {"A": 12, "B": 18, "C": 22, "D": 29, "E": 32})
    for row_index in range(2, 19):
        for column in (2, 3, 4):
            retained_sheet.cell(row_index, column).number_format = '$#,##0'
        retained_sheet.cell(row_index, 5).number_format = "0.00%"
    retained_sheet.conditional_formatting.add("D3:D18", FormulaRule(formula=["D3<0"], fill=PatternFill("solid", fgColor=PALE_RED)))
    add_native_table(retained_sheet, "RetainedRevenueTable")
    retained_sheet.print_area = "A1:E18"

    removed_sheet = workbook.create_sheet("Removed_Years")
    removed_sheet.append(["Year", "Revenue_USD", "Shortfall_to_Threshold"])
    for row_index, (year, revenue) in enumerate(removed, start=2):
        removed_sheet.append([year, revenue, f"={THRESHOLD}-B{row_index}"])
    style_sheet(removed_sheet, {"A": 14, "B": 20, "C": 25}, orientation="portrait")
    for row_index in range(2, 5):
        for column in (2, 3):
            removed_sheet.cell(row_index, column).number_format = '$#,##0'
    add_native_table(removed_sheet, "RemovedRevenueTable")
    removed_sheet.print_area = "A1:C4"

    summary = workbook.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    metrics = [
        ("Original_Years", "=COUNTA(Source_Data!A2:A21)"),
        ("Retained_Years", "=COUNTA(Retained_Years!A2:A18)"),
        ("Removed_Years", "=COUNTA(Removed_Years!A2:A4)"),
        ("Retention_Rate", "=B3/B2"),
        ("Retained_Total", "=SUM(Retained_Years!B2:B18)"),
        ("Retained_Average", "=AVERAGE(Retained_Years!B2:B18)"),
        ("Highest_Revenue_Year", "=INDEX(Retained_Years!A2:A18,MATCH(MAX(Retained_Years!B2:B18),Retained_Years!B2:B18,0))"),
        ("Highest_Revenue", "=MAX(Retained_Years!B2:B18)"),
        ("Lowest_Retained_Year", "=INDEX(Retained_Years!A2:A18,MATCH(MIN(Retained_Years!B2:B18),Retained_Years!B2:B18,0))"),
        ("Lowest_Retained_Revenue", "=MIN(Retained_Years!B2:B18)"),
        ("First_to_Last_Change", "=Retained_Years!B18-Retained_Years!B2"),
        ("First_to_Last_Change_Pct", "=Retained_Years!B18/Retained_Years!B2-1"),
    ]
    for label, formula in metrics:
        summary.append([label, formula])
    style_sheet(summary, {"A": 34, "B": 23}, orientation="landscape")
    for row_index in range(2, 14):
        summary.cell(row_index, 1).fill = PatternFill("solid", fgColor=PALE_BLUE)
        summary.cell(row_index, 1).font = Font(name="Aptos", size=10, bold=True, color=NAVY)
        summary.cell(row_index, 2).font = Font(name="Aptos Display", size=12, bold=True, color=TEAL)
    summary.cell(5, 2).number_format = "0.00%"
    for row_index in (6, 7, 9, 11, 12):
        summary.cell(row_index, 2).number_format = '$#,##0.00'
    summary.cell(13, 2).number_format = "0.00%"
    summary.conditional_formatting.add("B13", FormulaRule(formula=["B13<0"], fill=PatternFill("solid", fgColor=PALE_GOLD)))
    add_native_table(summary, "RevenueSummaryTable")
    chart = LineChart()
    chart.style = 13
    chart.title = "Retained revenue by year"
    chart.y_axis.title = "Revenue (USD)"
    chart.x_axis.title = "Year"
    chart.height = 7.2
    chart.width = 13.0
    chart.legend = None
    chart.add_data(Reference(retained_sheet, min_col=2, min_row=1, max_row=18), titles_from_data=True)
    chart.set_categories(Reference(retained_sheet, min_col=1, min_row=2, max_row=18))
    summary.add_chart(chart, "D2")
    summary.print_area = "A1:M13"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    recalculate_workbook(output_path)
    restore_frozen_header_rows(output_path, sheet_count=5)
    normalize_ooxml(output_path)


def recalculate_workbook(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="revenue-xlsx-recalc-") as tmp:
        root = Path(tmp)
        input_root = root / "input"
        output_root = root / "output"
        profile = root / "profile"
        input_root.mkdir()
        output_root.mkdir()
        profile.mkdir()
        source = input_root / path.name
        shutil.copy2(path, source)
        proc = subprocess.run(
            ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "xlsx", "--outdir", str(output_root), str(source)],
            stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300,
        )
        converted = output_root / path.name
        if proc.returncode != 0 or not converted.is_file():
            raise RuntimeError(f"LibreOffice 公式重算失败：{proc.stdout[-1000:]}")
        shutil.copy2(converted, path)


def restore_frozen_header_rows(path: Path, sheet_count: int) -> None:
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with tempfile.TemporaryDirectory(prefix="revenue-xlsx-freeze-") as tmp:
        target = Path(tmp) / path.name
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w") as destination:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename in {f"xl/worksheets/sheet{index}.xml" for index in range(1, sheet_count + 1)}:
                    root = etree.fromstring(data)
                    sheet_view = root.find(f"{{{namespace}}}sheetViews/{{{namespace}}}sheetView")
                    if sheet_view is not None:
                        for existing in sheet_view.findall(f"{{{namespace}}}pane"):
                            sheet_view.remove(existing)
                        pane = etree.Element(f"{{{namespace}}}pane")
                        pane.set("ySplit", "1")
                        pane.set("topLeftCell", "A2")
                        pane.set("activePane", "bottomLeft")
                        pane.set("state", "frozen")
                        sheet_view.insert(0, pane)
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                destination.writestr(info, data)
        shutil.copy2(target, path)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE, size: str = "6") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def add_word_table(document: Document, headers: list[str], data: list[list[str]], widths: list[float]):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, NAVY)
        set_cell_border(cell, NAVY)
        cell.width = Cm(widths[index])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(8)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(data):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.text = str(value)
            set_cell_shading(cell, WHITE if row_index % 2 == 0 else "F4F7FA")
            set_cell_border(cell)
            cell.width = Cm(widths[index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(8)
                    run.font.color.rgb = RGBColor.from_string(NAVY)
    return table


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Revenue screening  |  ")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)


def build_report(rows: list[tuple[int, int]], output_path: Path) -> None:
    retained = [(year, revenue) for year, revenue in rows if revenue >= THRESHOLD]
    removed = [(year, revenue) for year, revenue in rows if revenue < THRESHOLD]
    retained_total = sum(revenue for _, revenue in retained)
    retained_average = Decimal(retained_total) / Decimal(len(retained))
    highest_year, highest_revenue = max(retained, key=lambda item: item[1])
    lowest_year, lowest_revenue = min(retained, key=lambda item: item[1])
    first_year, first_revenue = retained[0]
    last_year, last_revenue = retained[-1]
    first_last_change = last_revenue - first_revenue
    first_last_pct = Decimal(last_revenue) / Decimal(first_revenue) - Decimal("1")

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.55)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    document.core_properties.creator = "EvoCowork OfficeBench adapter"
    document.core_properties.title = "Revenue Threshold Screening Report"
    document.core_properties.created = FIXED_DT.replace(tzinfo=None)
    document.core_properties.modified = FIXED_DT.replace(tzinfo=None)
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(10)
    document.styles["Normal"].font.color.rgb = RGBColor.from_string(NAVY)
    for name in ("Title", "Heading 1", "Heading 2"):
        document.styles[name].font.name = "Aptos Display"
        document.styles[name].font.color.rgb = RGBColor.from_string(NAVY)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title.add_run("Revenue Threshold Screening Report")
    title_run.font.size = Pt(27)
    title_run.font.bold = True
    subtitle = document.add_paragraph()
    subtitle_run = subtitle.add_run("2004–2023 revenue series  |  Threshold: $2.0 million  |  USD")
    subtitle_run.font.name = "Aptos"
    subtitle_run.font.size = Pt(10)
    subtitle_run.font.color.rgb = RGBColor.from_string(TEAL)

    add_word_table(
        document,
        ["Original years", "Retained", "Removed", "Retention rate"],
        [["20", "17", "3", "85.00%"]],
        [4.5, 4.5, 4.5, 4.5],
    )

    document.add_heading("Screening Rule", level=1)
    document.add_paragraph(
        "Years with revenue greater than or equal to $2,000,000 are retained. Years strictly below $2,000,000 are removed. The equality boundary is therefore included in the retained set."
    )

    document.add_heading("Executive Summary", level=1)
    document.add_paragraph(
        f"The source contains 20 annual observations. The screening retains 17 years and removes 3, producing an 85.00% retention rate. "
        f"Retained revenue totals {money(retained_total)}, with an average of {money(retained_average)}."
    )
    add_word_table(
        document,
        ["Metric", "Result"],
        [
            ["Retained total", money(retained_total)],
            ["Retained average", money(retained_average)],
            ["Highest retained year", f"{highest_year} — {money(highest_revenue)}"],
            ["Lowest retained year", f"{lowest_year} — {money(lowest_revenue)}"],
        ],
        [8.5, 9.5],
    )

    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_heading("Removed Years", level=1)
    removed_rows = [
        [str(year), money(revenue), money(THRESHOLD - revenue)]
        for year, revenue in removed
    ]
    add_word_table(document, ["Year", "Revenue", "Shortfall to threshold"], removed_rows, [4.0, 6.5, 7.5])
    document.add_paragraph(
        f"The three removed years are 2006, 2009 and 2017. Their combined shortfall to the threshold is {money(sum(THRESHOLD - revenue for _, revenue in removed))}."
    )

    document.add_heading("Retained Revenue Findings", level=1)
    document.add_paragraph(
        f"The highest retained revenue is {money(highest_revenue)} in {highest_year}. The lowest retained revenue is {money(lowest_revenue)} in {lowest_year}."
    )
    document.add_paragraph(
        f"From the first retained observation ({first_year}, {money(first_revenue)}) to the last retained observation ({last_year}, {money(last_revenue)}), "
        f"revenue changes by {money(first_last_change)} ({first_last_pct:.2%}). This is an endpoint comparison, not a claim of a continuous trend."
    )
    top_three = sorted(retained, key=lambda item: item[1], reverse=True)[:3]
    add_word_table(
        document,
        ["Rank", "Year", "Revenue"],
        [[str(index), str(year), money(revenue)] for index, (year, revenue) in enumerate(top_three, start=1)],
        [4.0, 5.5, 8.5],
    )

    document.add_heading("Evidence Boundary", level=1)
    document.add_paragraph(
        "All calculations use only company_revenues.xlsx. The source supplies annual revenue values but no information about inflation, market conditions, products, costs, profit, management actions or forecasts. The report therefore describes the observed values and threshold result without assigning causes or predicting future performance."
    )

    add_page_number(document.sections[0].footer.paragraphs[0])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    normalize_ooxml(output_path)


def normalize_ooxml(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="revenue-ooxml-normalize-") as tmp:
        target = Path(tmp) / path.name
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w") as destination:
            for info in sorted(source.infolist(), key=lambda item: item.filename):
                data = source.read(info.filename)
                if info.filename == "docProps/core.xml":
                    root = etree.fromstring(data)
                    dcterms = "http://purl.org/dc/terms/"
                    for field in ("created", "modified"):
                        node = root.find(f"{{{dcterms}}}{field}")
                        if node is not None:
                            node.text = "2026-09-01T05:45:00Z"
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                elif info.filename.startswith("xl/charts/") and info.filename.endswith(".xml"):
                    root = etree.fromstring(data)
                    chart_ns = "http://schemas.openxmlformats.org/drawingml/2006/chart"
                    category_axis = root.find(f".//{{{chart_ns}}}catAx/{{{chart_ns}}}axId")
                    value_axis = root.find(f".//{{{chart_ns}}}valAx/{{{chart_ns}}}axId")
                    if category_axis is not None and value_axis is not None:
                        mapping = {category_axis.get("val", ""): "12017001", value_axis.get("val", ""): "12017002"}
                        for node in root.iter():
                            if etree.QName(node).localname in {"axId", "crossAx"} and node.get("val") in mapping:
                                node.set("val", mapping[node.get("val")])
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                copied = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
                copied.compress_type = zipfile.ZIP_DEFLATED
                copied.external_attr = info.external_attr
                copied.create_system = info.create_system
                copied.flag_bits = info.flag_bits
                destination.writestr(copied, data)
        shutil.copy2(target, path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for child in args.output_dir.iterdir():
        if child.is_file():
            child.unlink()
        elif child.is_dir():
            shutil.rmtree(child)
    rows = read_source(args.input_dir)
    build_workbook(rows, args.output_dir / XLSX_NAME)
    build_report(rows, args.output_dir / DOCX_NAME)


if __name__ == "__main__":
    main()
