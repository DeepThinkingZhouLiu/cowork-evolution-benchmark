#!/usr/bin/env python3
"""OfficeBench 3-8 收入门槛交付包的 v2 任务专属原子评分。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from openpyxl import load_workbook


TASK_ID="cowork-evo/officebench-3-8-revenue-threshold-report-package"
RUBRIC_VERSION="officebench-3-8-revenue-threshold-report-v0.2"
JUDGE_VERSION="revenue-threshold-report-atomic-v2.0"
XLSX="Revenue_Filtered_Analysis.xlsx";DOCX="Revenue_Threshold_Report.docx";REQUIRED=(XLSX,DOCX)
SOURCE_HASH="b7a5d913221fafb9f545e9367df2bcf07e665b93baefe726788a88532c937e28"
THRESHOLD=2_000_000
WEIGHTS={"R001":6,"R002":10,"R003":10,"R004":8,"R005":13,"R006":12,"R007":6,"R008":7,"R009":8,"R010":7,"R011":7,"R012":6}
PENALTIES={"P001":-12,"P002":-4}


def _load_v1():
    path=Path(__file__).with_name("judge_v1.py");spec=importlib.util.spec_from_file_location("revenue_judge_v1",path)
    if spec is None or spec.loader is None:raise RuntimeError("无法载入 v1 Judge")
    module=importlib.util.module_from_spec(spec);sys.modules[spec.name]=module;spec.loader.exec_module(module);return module


V1=_load_v1()
ROWS,TRUTH,SOURCE_ERRORS=V1.load_truth()
RETAINED=TRUTH.get("retained",[]);REMOVED=TRUTH.get("removed",[])


def _sha256(path:Path)->str:
    digest=hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda:handle.read(1024*1024),b""):digest.update(block)
    return digest.hexdigest()


def _task_root()->Path:return Path(__file__).resolve().parents[1]


def _source()->Path:
    harbor=Path("/app/input_files/company_revenues.xlsx")
    return harbor if harbor.is_file() else _task_root()/"data"/"input_files"/"company_revenues.xlsx"


def _norm(value:Any)->str:return re.sub(r"\s+"," ",str(value or "").replace("\u00a0"," ").replace("\v"," ")).strip().casefold()


def _same(actual:Any,expected:Any,tolerance:float=1e-7)->bool:
    if isinstance(expected,(int,float)) and not isinstance(expected,bool):
        try:return abs(float(actual)-float(expected))<=tolerance
        except (TypeError,ValueError):return False
    return _norm(actual)==_norm(expected)


def _formula(value:Any,*tokens:str)->bool:
    if not isinstance(value,str) or not value.startswith("="):return False
    normalized=re.sub(r"[\s'$]","",value).upper()
    return all(re.sub(r"[\s'$]","",token).upper() in normalized for token in tokens)


def _ratio(items:Iterable[tuple[bool,str]])->tuple[float,list[str],int,int]:
    checks=list(items);earned=sum(bool(ok) for ok,_ in checks)
    return earned/max(1,len(checks)),[label for ok,label in checks if not ok],earned,len(checks)


def _criterion(identifier:str,items:Iterable[tuple[bool,str]])->dict[str,Any]:
    score,failures,earned,total=_ratio(items)
    return {"criterion_id":identifier,"score":round(score,6),"evidence":"全部任务专属原子检查通过" if not failures else "未通过："+"；".join(failures[:24]),"method":"deterministic","raw":{"earned_units":earned,"total_units":total,"failure_count":len(failures)}}


def _core(artifact:Path)->tuple[bool,list[str],list[str]]:
    errors=[];hashes=[]
    if not artifact.is_dir():return False,["输出目录不存在"],hashes
    for name in REQUIRED:
        path=artifact/name
        if not path.is_file() or path.stat().st_size==0:errors.append(f"核心交付缺失或为空：{name}");continue
        try:
            hashes.append(_sha256(path))
            if not zipfile.is_zipfile(path):raise ValueError("不是有效 OOXML ZIP")
            if name==XLSX:
                book=load_workbook(path,read_only=True,data_only=False)
                if not book.sheetnames:raise ValueError("Excel 无工作表")
                book.close()
            else:
                document=Document(path)
                if not document.paragraphs and not document.tables:raise ValueError("Word 无内容")
        except Exception as exc:errors.append(f"核心交付无法解析：{name}: {type(exc).__name__}: {exc}")
    source=_source()
    if not source.is_file() or _sha256(source)!=SOURCE_HASH:errors.append("冻结 company_revenues.xlsx 缺失或哈希错误")
    return not errors,errors,hashes


def _fatal(errors:list[str],hashes:list[str])->dict[str,Any]:
    results=[_criterion(key,[(False,error) for error in errors]) for key in WEIGHTS]
    results.extend(_criterion(key,[(True,"未触发")]) for key in PENALTIES)
    return {"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":hashes or ["0"*64],"criterion_results":results,"reward":0.0,"passed":False,"hurdle_failures":["R001"],"penalties":[]}


def _recalculate(path:Path)->tuple[dict[str,list[list[Any]]]|None,str]:
    with tempfile.TemporaryDirectory(prefix="revenue-v2-recalc-") as raw:
        root=Path(raw);source=root/"source";output=root/"output";profile=root/"profile";source.mkdir();output.mkdir();copied=source/path.name;shutil.copy2(path,copied)
        env=os.environ.copy();env["HOME"]=str(root/"home");env["NO_PROXY"]="localhost,127.0.0.1,::1"
        try:completed=subprocess.run(["soffice","--headless",f"-env:UserInstallation=file://{profile}","--convert-to","xlsx","--outdir",str(output),str(copied)],stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,timeout=300,check=False,env=env)
        except Exception as exc:return None,str(exc)
        converted=output/path.name
        if completed.returncode!=0 or not converted.is_file():return None,f"退出={completed.returncode}；日志={completed.stdout[-400:]}"
        try:
            book=load_workbook(converted,data_only=True);values={sheet.title:[[cell.value for cell in row] for row in sheet.iter_rows()] for sheet in book.worksheets};book.close();return values,completed.stdout[-400:]
        except Exception as exc:return None,str(exc)


def _value(values:dict[str,list[list[Any]]]|None,sheet:str,row:int,column:int)->Any:
    try:return values[sheet][row-1][column-1] if values is not None else None
    except (KeyError,IndexError):return None


def _workbook(artifact:Path)->tuple[dict[str,list[tuple[bool,str]]],dict[str,Any],list[Any]]:
    findings={key:[] for key in ("R002","R003","R004","R005","R006","R007")};context={};danger=[];path=artifact/XLSX
    try:
        package=V1.inspect_package(path);danger=[*package.get("macros",[]),*package.get("external",[]),*package.get("missing_targets",[])]
        book=load_workbook(path,data_only=False);values,recalc_log=_recalculate(path)
    except Exception as exc:
        for key in findings:findings[key].append((False,f"Excel 检查失败：{exc}"))
        return findings,context,danger
    expected_sheets=V1.SHEETS
    for index,name in enumerate(expected_sheets):findings["R007"].append((index<len(book.sheetnames) and book.sheetnames[index]==name,f"工作表第 {index+1} 项错误"))

    source=book["Source_Data"] if "Source_Data" in book.sheetnames else None
    audit=book["Threshold_Audit"] if "Threshold_Audit" in book.sheetnames else None
    for column,expected in enumerate(V1.SOURCE_HEADERS,1):findings["R002"].append((source is not None and _same(source.cell(1,column).value,expected),f"Source_Data 表头第 {column} 列错误"))
    for column,expected in enumerate(V1.AUDIT_HEADERS,1):findings["R002"].append((audit is not None and _same(audit.cell(1,column).value,expected),f"Threshold_Audit 表头第 {column} 列错误"))
    audit_context=[]
    for row,(year,revenue) in enumerate(ROWS,2):
        findings["R002"].extend([
            (_same(source.cell(row,1).value if source else None,year),f"Source_Data!A{row} 年份错误"),
            (_same(source.cell(row,2).value if source else None,revenue),f"Source_Data!B{row} 收入错误"),
            (_same(audit.cell(row,1).value if audit else None,year),f"Threshold_Audit!A{row} 年份错误"),
            (_same(audit.cell(row,2).value if audit else None,revenue),f"Threshold_Audit!B{row} 收入错误"),
            (_same(audit.cell(row,3).value if audit else None,THRESHOLD),f"Threshold_Audit!C{row} 门槛错误"),
        ])
        expected_status="Retain" if revenue>=THRESHOLD else "Remove";expected_rank=TRUTH["rank"][year]
        checks=((4,(f"B{row}>=C{row}",),expected_status),(5,(f"B{row}-C{row}",),revenue-THRESHOLD),(6,("RANK",f"B{row}","B2:B21"),expected_rank))
        for column,tokens,target in checks:
            findings["R005"].extend([(_formula(audit.cell(row,column).value if audit else None,*tokens),f"Threshold_Audit!{chr(64+column)}{row} 公式错误"),(_same(_value(values,"Threshold_Audit",row,column),target),f"Threshold_Audit!{chr(64+column)}{row} 重算错误：{recalc_log}")])
        audit_context.append((year,revenue,_value(values,"Threshold_Audit",row,4)))
    context["audit"]=audit_context

    retained=book["Retained_Years"] if "Retained_Years" in book.sheetnames else None
    for column,expected in enumerate(V1.RETAINED_HEADERS,1):findings["R003"].append((retained is not None and _same(retained.cell(1,column).value,expected),f"Retained_Years 表头第 {column} 列错误"))
    retained_context=[]
    for row,(year,revenue) in enumerate(RETAINED,2):
        findings["R003"].extend([(_same(retained.cell(row,1).value if retained else None,year),f"Retained_Years!A{row} 错误"),(_same(retained.cell(row,2).value if retained else None,revenue),f"Retained_Years!B{row} 错误")])
        findings["R005"].extend([(_formula(retained.cell(row,3).value if retained else None,f"B{row}-2000000"),f"Retained_Years!C{row} 公式错误"),(_same(_value(values,"Retained_Years",row,3),revenue-THRESHOLD),f"Retained_Years!C{row} 重算错误")])
        previous=None if row==2 else RETAINED[row-3][1]
        if previous is None:
            findings["R005"].extend([(retained is not None and retained.cell(row,4).value is None,"Retained_Years!D2 应为空"),(retained is not None and retained.cell(row,5).value is None,"Retained_Years!E2 应为空")])
        else:
            findings["R005"].extend([
                (_formula(retained.cell(row,4).value if retained else None,f"B{row}-B{row-1}"),f"Retained_Years!D{row} 公式错误"),
                (_same(_value(values,"Retained_Years",row,4),revenue-previous),f"Retained_Years!D{row} 重算错误"),
                (_formula(retained.cell(row,5).value if retained else None,"IFERROR",f"B{row}/B{row-1}-1"),f"Retained_Years!E{row} 公式错误"),
                (_same(_value(values,"Retained_Years",row,5),revenue/previous-1),f"Retained_Years!E{row} 重算错误"),
            ])
        retained_context.append((year,revenue))
    context["retained"]=retained_context

    removed=book["Removed_Years"] if "Removed_Years" in book.sheetnames else None
    for column,expected in enumerate(V1.REMOVED_HEADERS,1):findings["R004"].append((removed is not None and _same(removed.cell(1,column).value,expected),f"Removed_Years 表头第 {column} 列错误"))
    removed_context=[]
    for row,(year,revenue) in enumerate(REMOVED,2):
        findings["R004"].extend([(_same(removed.cell(row,1).value if removed else None,year),f"Removed_Years!A{row} 年份错误"),(_same(removed.cell(row,2).value if removed else None,revenue),f"Removed_Years!B{row} 收入错误")])
        findings["R005"].extend([(_formula(removed.cell(row,3).value if removed else None,f"2000000-B{row}"),f"Removed_Years!C{row} 公式错误"),(_same(_value(values,"Removed_Years",row,3),THRESHOLD-revenue),f"Removed_Years!C{row} 重算错误")])
        removed_context.append((year,revenue,_value(values,"Removed_Years",row,3)))
    context["removed"]=removed_context

    summary=book["Summary"] if "Summary" in book.sheetnames else None
    expected_values=[20,17,3,.85,TRUTH["retained_total"],float(TRUTH["retained_average"]),TRUTH["highest_year"],TRUTH["highest_revenue"],TRUTH["lowest_year"],TRUTH["lowest_revenue"],TRUTH["first_last_change"],float(TRUTH["first_last_pct"])]
    formula_tokens=[("COUNTA","SOURCE_DATA!A2:A21"),("COUNTA","RETAINED_YEARS!A2:A18"),("COUNTA","REMOVED_YEARS!A2:A4"),("B3/B2",),("SUM","RETAINED_YEARS!B2:B18"),("AVERAGE","RETAINED_YEARS!B2:B18"),("INDEX","MATCH","MAX","RETAINED_YEARS!B2:B18"),("MAX","RETAINED_YEARS!B2:B18"),("INDEX","MATCH","MIN","RETAINED_YEARS!B2:B18"),("MIN","RETAINED_YEARS!B2:B18"),("RETAINED_YEARS!B18-RETAINED_YEARS!B2",),("RETAINED_YEARS!B18/RETAINED_YEARS!B2-1",)]
    summary_context={}
    for row,(label,target,tokens) in enumerate(zip(V1.SUMMARY_LABELS,expected_values,formula_tokens,strict=True),2):
        findings["R006"].extend([(_same(summary.cell(row,1).value if summary else None,label),f"Summary!A{row} 标签错误"),(_formula(summary.cell(row,2).value if summary else None,*tokens),f"Summary!B{row} 公式错误"),(_same(_value(values,"Summary",row,2),target),f"Summary!B{row} 重算错误")]);summary_context[label]=_value(values,"Summary",row,2)
    context["summary"]=summary_context

    expected_tables={"Source_Data":"RevenueSourceTable","Threshold_Audit":"RevenueThresholdAuditTable","Retained_Years":"RetainedRevenueTable","Removed_Years":"RemovedRevenueTable","Summary":"RevenueSummaryTable"}
    for sheet_name,table_name in expected_tables.items():
        sheet=book[sheet_name] if sheet_name in book.sheetnames else None
        findings["R007"].extend([(sheet is not None and set(sheet.tables)=={table_name},f"{sheet_name} 原生表错误"),(sheet is not None and sheet.freeze_panes=="A2",f"{sheet_name} 未冻结首行"),(sheet is not None and bool(sheet.auto_filter.ref),f"{sheet_name} 缺筛选")])
    findings["R007"].extend([(summary is not None and len(summary._charts)==1,"Summary 缺可编辑折线图"),(sum(len(sheet.conditional_formatting) for sheet in book.worksheets)>=3,"条件格式少于 3 处"),(not any(len(sheet._images) for sheet in book.worksheets),"Excel 用图片代替可编辑数据"),(not danger,f"OOXML 含危险/断裂对象：{danger}")])
    book.close();return findings,context,danger


def _table_rows(table)->list[list[str]]:return [[cell.text for cell in row.cells] for row in table.rows]


def _word(artifact:Path)->tuple[dict[str,list[tuple[bool,str]]],dict[str,Any],list[str]]:
    findings={key:[] for key in ("R008","R009","R010")};document=Document(artifact/DOCX);text=V1.word_text(document);lower=_norm(text);headings=[V1.normalize(p.text) for p in document.paragraphs if str(p.style.name).startswith("Heading")];unsupported=V1.unsupported_claims(text)
    findings["R008"].append(("Revenue Threshold Screening Report" in text,"Word 标题错误"))
    for heading in ("Screening Rule","Executive Summary","Removed Years","Retained Revenue Findings","Evidence Boundary"):findings["R008"].append((heading in headings,f"Word 缺章节 {heading}"))
    for token in ("greater than or equal to $2,000,000","strictly below $2,000,000","equality boundary"):findings["R008"].append((token.casefold() in lower,f"Word 门槛规则缺 {token}"))
    findings["R008"].extend([(len(document.tables)==4,f"Word 原生表数错误：{len(document.tables)}"),(not document.inline_shapes,"Word 用图片代替表格")])

    expected_tables=[
        [["Original years","Retained","Removed","Retention rate"],["20","17","3","85.00%"]],
        [["Metric","Result"],["Retained total","$56,172,635.00"],["Retained average","$3,304,272.65"],["Highest retained year","2014 — $4,983,791.00"],["Lowest retained year","2021 — $2,005,820.00"]],
        [["Year","Revenue","Shortfall to threshold"],["2006","$1,193,225.00","$806,775.00"],["2009","$1,791,112.00","$208,888.00"],["2017","$1,463,500.00","$536,500.00"]],
        [["Rank","Year","Revenue"],["1","2014","$4,983,791.00"],["2","2015","$4,611,210.00"],["3","2004","$4,439,044.00"]],
    ]
    table_rows=[_table_rows(table) for table in document.tables]
    for table_index in (0,2):
        expected=expected_tables[table_index];table=document.tables[table_index] if table_index<len(document.tables) else None
        for row_index,row in enumerate(expected):
            for column_index,target in enumerate(row):
                actual=table.rows[row_index].cells[column_index].text if table is not None and row_index<len(table.rows) and column_index<len(table.rows[row_index].cells) else None
                findings["R009"].append((_same(actual,target),f"Word 表{table_index+1}第{row_index+1}行第{column_index+1}列错误"))
    for token in ("20 annual observations","17 years","removes 3","85.00%","$1,552,163.00"):findings["R009"].append((token.casefold() in lower,f"Word 缺筛选事实 {token}"))

    for table_index in (1,3):
        expected=expected_tables[table_index];table=document.tables[table_index] if table_index<len(document.tables) else None
        for row_index,row in enumerate(expected):
            for column_index,target in enumerate(row):
                actual=table.rows[row_index].cells[column_index].text if table is not None and row_index<len(table.rows) and column_index<len(table.rows[row_index].cells) else None
                findings["R010"].append((_same(actual,target),f"Word 表{table_index+1}第{row_index+1}行第{column_index+1}列错误"))
    for token in ("$56,172,635.00","$3,304,272.65","$4,983,791.00 in 2014","$2,005,820.00 in 2021","-$1,606,499.00","-36.19%","endpoint comparison","not a claim of a continuous trend","use only company_revenues.xlsx","without assigning causes","predicting future performance"):findings["R010"].append((token.casefold() in lower,f"Word 缺经营事实/边界 {token}"))
    return findings,{"text":text,"tables":table_rows},unsupported


def evaluate(artifact:Path,scoring_version:str|None=None)->dict[str,Any]:
    if (scoring_version or os.environ.get("SCORING_VERSION","v2")).casefold()=="v1":return V1.evaluate(artifact)
    core_ok,errors,hashes=_core(artifact)
    if not core_ok:return _fatal(errors,hashes)
    criteria={key:[] for key in WEIGHTS};criteria["R001"].append((not SOURCE_ERRORS,"两份核心交付和冻结输入可解析"))
    workbook,workbook_context,danger=_workbook(artifact)
    for key,items in workbook.items():criteria[key].extend(items)
    word,word_context,unsupported=_word(artifact)
    for key,items in word.items():criteria[key].extend(items)
    summary=workbook_context.get("summary",{});text=word_context.get("text","").casefold()
    cross_tokens={"Original_Years":"20 annual observations","Retained_Years":"17 years","Removed_Years":"removes 3","Retained_Total":"$56,172,635.00","Highest_Revenue_Year":"2014","Lowest_Retained_Year":"2021"}
    for label,token in cross_tokens.items():criteria["R011"].append((_same(summary.get(label),dict(zip(V1.SUMMARY_LABELS,[20,17,3,.85,TRUTH["retained_total"],float(TRUTH["retained_average"]),TRUTH["highest_year"],TRUTH["highest_revenue"],TRUTH["lowest_year"],TRUTH["lowest_revenue"],TRUTH["first_last_change"],float(TRUTH["first_last_pct"])]))[label]) and token.casefold() in text,f"Excel/Word 一致性 {label} 错误"))
    render=V1.render_office(artifact);criteria["R012"].extend([(not render.get("errors"),f"LibreOffice 渲染错误：{render.get('errors')}"),(render.get("xlsx_pages")==5,f"Excel 页数错误：{render.get('xlsx_pages')}"),(render.get("docx_pages")==2,f"Word 页数错误：{render.get('docx_pages')}")])
    extra=sorted(path.name for path in artifact.iterdir() if path.name not in REQUIRED);suspicious=V1.suspicious_content(artifact);p1_score=min(1.0,len(unsupported)/2.0);p2_bad=bool(extra or danger or suspicious)
    positive=[_criterion(key,criteria[key]) for key in WEIGHTS];base=sum(WEIGHTS[item["criterion_id"]]*float(item["score"]) for item in positive)/100.0;applied=[]
    if p1_score:applied.append({"id":"P001","amount":round(.12*p1_score,6),"reason":"无来源原因/利润等义/预测"})
    if p2_bad:applied.append({"id":"P002","amount":.04,"reason":f"额外交付或危险对象：extra={extra}"})
    reward=round(max(0.0,base-sum(item["amount"] for item in applied)),6)
    positive.extend([{"criterion_id":"P001","score":round(p1_score,6),"evidence":"未触发" if not p1_score else "检测到无来源断言","method":"deterministic","raw":{"unsupported":unsupported}},{"criterion_id":"P002","score":float(p2_bad),"evidence":"未触发" if not p2_bad else "检测到额外交付/危险对象","method":"deterministic","raw":{"extra":extra,"danger":danger,"suspicious":suspicious}}])
    return {"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":hashes,"criterion_results":positive,"base_score":round(base,6),"reward":reward,"passed":reward>=.8,"hurdle_failures":[],"penalties":applied}


def main()->None:
    parser=argparse.ArgumentParser();parser.add_argument("--artifact-dir",type=Path,required=True);parser.add_argument("--result",type=Path,required=True);parser.add_argument("--reward-file",type=Path,required=True);parser.add_argument("--scoring-version",choices=("v1","v2"),default=os.environ.get("SCORING_VERSION","v2"));args=parser.parse_args()
    try:result=evaluate(args.artifact_dir,args.scoring_version)
    except Exception as exc:result={"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"reward":0.0,"passed":False,"fatal_error":f"{type(exc).__name__}: {exc}","hurdle_failures":["R001"]}
    args.result.parent.mkdir(parents=True,exist_ok=True);args.reward_file.parent.mkdir(parents=True,exist_ok=True);args.result.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n",encoding="utf-8");args.reward_file.write_text(f"{float(result['reward']):.6f}\n",encoding="utf-8");print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=="__main__":main()
