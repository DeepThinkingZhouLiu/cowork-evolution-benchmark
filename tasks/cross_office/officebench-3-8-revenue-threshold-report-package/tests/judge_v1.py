#!/usr/bin/env python3
"""OfficeBench 3-8 收入筛选题的来源重算、公式与跨文件一致性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from docx import Document
from lxml import etree
from openpyxl import load_workbook


TASK_ID = "cowork-evo/officebench-3-8-revenue-threshold-report-package"
RUBRIC_VERSION = "officebench-3-8-revenue-threshold-report-v0.1"
JUDGE_VERSION = "officebench-3-8-revenue-threshold-report-deterministic-v0.1"
XLSX_NAME = "Revenue_Filtered_Analysis.xlsx"
DOCX_NAME = "Revenue_Threshold_Report.docx"
REQUIRED_FILES = [XLSX_NAME, DOCX_NAME]
SOURCE_HASH = "b7a5d913221fafb9f545e9367df2bcf07e665b93baefe726788a88532c937e28"
THRESHOLD = 2_000_000
SHEETS = ["Source_Data", "Threshold_Audit", "Retained_Years", "Removed_Years", "Summary"]
SOURCE_HEADERS = ["Year", "Revenue_USD"]
AUDIT_HEADERS = ["Year", "Revenue_USD", "Threshold_USD", "Status", "Margin_to_Threshold", "Revenue_Rank"]
RETAINED_HEADERS = ["Year", "Revenue_USD", "Margin_to_Threshold", "Change_vs_Previous_Retained", "Change_Pct_vs_Previous_Retained"]
REMOVED_HEADERS = ["Year", "Revenue_USD", "Shortfall_to_Threshold"]
SUMMARY_LABELS = [
    "Original_Years", "Retained_Years", "Removed_Years", "Retention_Rate", "Retained_Total",
    "Retained_Average", "Highest_Revenue_Year", "Highest_Revenue", "Lowest_Retained_Year",
    "Lowest_Retained_Revenue", "First_to_Last_Change", "First_to_Last_Change_Pct",
]
SUSPICIOUS_MARKERS = ["solution/golden", "judge.py", "expectations.json", "judge-result", "/root/", "/home/", "c:\\"]

SPECS = {
    "R001": ("完整交付物与冻结输入", 5, "hurdle", 0.0),
    "R002": ("二十年来源与审计数据", 8, "hurdle", 0.3),
    "R003": ("十七个保留年份", 10, "hurdle", 0.35),
    "R004": ("三个移除年份与差额", 8, "hurdle", 0.4),
    "R005": ("筛选与变化公式", 10, "hurdle", 0.4),
    "R006": ("十二项公式汇总", 12, "hurdle", 0.45),
    "R007": ("Excel 原生分析结构", 6, "positive", None),
    "R008": ("Word 结构与筛选规则", 7, "hurdle", 0.5),
    "R009": ("Word 保留与移除真值", 9, "hurdle", 0.5),
    "R010": ("Word 经营指标与证据边界", 7, "hurdle", 0.55),
    "R011": ("Excel 与 Word 一致", 10, "hurdle", 0.4),
    "R012": ("可渲染与专业排版", 8, "positive", None),
    "R013": ("无年份、金额或分类错误", -30, "penalty", None),
    "R014": ("无跨文件事实冲突", -25, "penalty", None),
    "R015": ("无缺失、损坏、静态化或调试泄漏", -25, "penalty", None),
    "R016": ("无输入不支持的原因与预测", -15, "penalty", None),
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def normalize(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\u00a0", " ").replace("\v", " ")).strip()


def approx(left: object, right: object, tolerance: float = 1e-7) -> bool:
    try:
        return abs(float(left) - float(right)) <= tolerance
    except (TypeError, ValueError):
        return False


def money(value: Decimal | int) -> str:
    rounded = Decimal(value).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return f"-${abs(rounded):,.2f}" if rounded < 0 else f"${rounded:,.2f}"


def source_path() -> Path | None:
    task_root = Path(__file__).resolve().parents[1]
    candidates = [
        Path("/app/input_files/company_revenues.xlsx"),
        task_root / "data/input_files/company_revenues.xlsx",
        task_root / "environment/input_files/company_revenues.xlsx",
    ]
    return next((path for path in candidates if path.is_file()), None)


def load_truth() -> tuple[list[tuple[int, int]], dict[str, object], list[str]]:
    errors: list[str] = []
    path = source_path()
    if path is None:
        return [], {}, ["source_missing"]
    actual_hash = sha256(path)
    if actual_hash != SOURCE_HASH:
        return [], {}, [f"source_sha256={actual_hash}"]
    try:
        workbook = load_workbook(path, data_only=False, read_only=True)
        sheet = workbook["Sheet1"]
        headers = [normalize(cell.value) for cell in sheet[1]]
        rows = [(int(row[0]), int(row[1])) for row in sheet.iter_rows(min_row=2, values_only=True) if any(value is not None for value in row)]
        workbook.close()
    except Exception as exc:
        return [], {}, [f"{type(exc).__name__}:{exc}"]
    if headers != ["Year", "Revenue ($)"] or len(rows) != 20 or [year for year, _ in rows] != list(range(2004, 2024)):
        errors.append(f"source_structure={headers}/{len(rows)}")
    retained = [(year, revenue) for year, revenue in rows if revenue >= THRESHOLD]
    removed = [(year, revenue) for year, revenue in rows if revenue < THRESHOLD]
    retained_total = sum(revenue for _, revenue in retained)
    retained_average = Decimal(retained_total) / Decimal(len(retained))
    highest_year, highest_revenue = max(retained, key=lambda item: item[1])
    lowest_year, lowest_revenue = min(retained, key=lambda item: item[1])
    first_last_change = retained[-1][1] - retained[0][1]
    truth = {
        "retained": retained,
        "removed": removed,
        "retained_total": retained_total,
        "retained_average": retained_average,
        "highest_year": highest_year,
        "highest_revenue": highest_revenue,
        "lowest_year": lowest_year,
        "lowest_revenue": lowest_revenue,
        "first_last_change": first_last_change,
        "first_last_pct": Decimal(retained[-1][1]) / Decimal(retained[0][1]) - Decimal("1"),
        "rank": {year: index for index, (year, _) in enumerate(sorted(rows, key=lambda item: item[1], reverse=True), start=1)},
    }
    return rows, truth, errors


def inspect_package(path: Path) -> dict[str, object]:
    result: dict[str, object] = {"broken": None, "missing_targets": [], "macros": [], "external": [], "media": []}
    if not path.is_file():
        result["broken"] = "missing"
        return result
    try:
        with zipfile.ZipFile(path) as package:
            names = set(package.namelist())
            result["macros"] = sorted(name for name in names if "vbaProject" in name or name.endswith(".bin"))
            result["media"] = sorted(name for name in names if "/media/" in name)
            for rel_name in [name for name in names if name.endswith(".rels")]:
                root = etree.fromstring(package.read(rel_name))
                for rel in root:
                    target = rel.get("Target", "")
                    if rel.get("TargetMode") == "External":
                        result["external"].append(target)
                        continue
                    if not target or target.startswith("#"):
                        continue
                    if target.startswith("/"):
                        resolved = target.lstrip("/")
                    else:
                        base = Path(rel_name).parent
                        if base.name == "_rels":
                            base = base.parent
                        resolved = str((base / target).as_posix())
                        while "/../" in resolved:
                            resolved = re.sub(r"[^/]+/\.\./", "", resolved, count=1)
                    if resolved not in names:
                        result["missing_targets"].append({"rels": rel_name, "target": target, "resolved": resolved})
    except Exception as exc:
        result["broken"] = f"{type(exc).__name__}: {exc}"
    return result


def recalculate_candidate(path: Path) -> tuple[Path | None, str]:
    temp_root = Path(tempfile.mkdtemp(prefix="revenue-judge-recalc-"))
    input_root = temp_root / "input"
    output_root = temp_root / "output"
    profile = temp_root / "profile"
    input_root.mkdir()
    output_root.mkdir()
    profile.mkdir()
    source = input_root / path.name
    shutil.copy2(path, source)
    command = [
        "soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "xlsx",
        "--outdir", str(output_root), str(source),
    ]
    try:
        proc = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300)
    except Exception as exc:
        return None, f"{type(exc).__name__}: {exc}"
    converted = output_root / path.name
    if proc.returncode != 0 or not converted.is_file():
        return None, proc.stdout[-1000:]
    return converted, proc.stdout[-1000:]


def headers(sheet) -> list[str]:
    return [normalize(cell.value) for cell in sheet[1]]


def parse_workbook(path: Path, rows: list[tuple[int, int]], truth: dict[str, object]) -> dict[str, object]:
    result: dict[str, object] = {
        "readable": False, "sheetnames": [], "source_errors": [], "retained_errors": [], "removed_errors": [],
        "formula_errors": [], "summary_errors": [], "native": {},
    }
    try:
        formulas = load_workbook(path, data_only=False)
    except Exception as exc:
        result["error"] = f"{type(exc).__name__}: {exc}"
        return result
    recalculated, log = recalculate_candidate(path)
    result["recalc_log"] = log
    if recalculated is None:
        return result
    try:
        values = load_workbook(recalculated, data_only=True)
    except Exception as exc:
        result["error"] = f"recalc load:{type(exc).__name__}:{exc}"
        return result
    result["readable"] = True
    result["sheetnames"] = formulas.sheetnames
    if formulas.sheetnames != SHEETS:
        result["source_errors"].append(f"sheetnames={formulas.sheetnames}")
        return result

    source = formulas["Source_Data"]
    if headers(source) != SOURCE_HEADERS or source.max_row != 21 or source.max_column != 2:
        result["source_errors"].append(f"source_structure={headers(source)}/{source.max_row}x{source.max_column}")
    for row_index, expected in enumerate(rows, start=2):
        got = (source.cell(row_index, 1).value, source.cell(row_index, 2).value)
        if not approx(got[0], expected[0], 0) or not approx(got[1], expected[1], 0):
            result["source_errors"].append(f"Source_Data!A{row_index}:B{row_index}={got}")

    audit_f = formulas["Threshold_Audit"]
    audit_v = values["Threshold_Audit"]
    if headers(audit_f) != AUDIT_HEADERS or audit_f.max_row != 21 or audit_f.max_column != 6:
        result["source_errors"].append(f"audit_structure={headers(audit_f)}/{audit_f.max_row}x{audit_f.max_column}")
    for row_index, (year, revenue) in enumerate(rows, start=2):
        expected_status = "Retain" if revenue >= THRESHOLD else "Remove"
        expected_values = [year, revenue, THRESHOLD, expected_status, revenue - THRESHOLD, truth["rank"][year]]
        for column, expected in enumerate(expected_values, start=1):
            got = audit_v.cell(row_index, column).value
            if column == 4:
                valid = normalize(got) == expected
            else:
                valid = approx(got, expected, 1e-7)
            if not valid:
                result["source_errors"].append(f"Threshold_Audit!{audit_v.cell(row_index, column).coordinate}={got!r}")
        for column in (4, 5, 6):
            if audit_f.cell(row_index, column).data_type != "f":
                result["formula_errors"].append(f"Threshold_Audit!{audit_f.cell(row_index, column).coordinate}:not_formula")
        formula_status = normalize(audit_f.cell(row_index, 4).value).upper().replace("$", "")
        formula_margin = normalize(audit_f.cell(row_index, 5).value).upper().replace("$", "")
        formula_rank = normalize(audit_f.cell(row_index, 6).value).upper().replace("$", "")
        if f"B{row_index}>=C{row_index}" not in formula_status:
            result["formula_errors"].append(f"D{row_index}:wrong_threshold")
        if f"B{row_index}-C{row_index}" not in formula_margin:
            result["formula_errors"].append(f"E{row_index}:wrong_source")
        if f"B{row_index}" not in formula_rank or "B2:B21" not in formula_rank or "RANK" not in formula_rank:
            result["formula_errors"].append(f"F{row_index}:wrong_source")

    retained_f = formulas["Retained_Years"]
    retained_v = values["Retained_Years"]
    expected_retained = truth["retained"]
    if headers(retained_f) != RETAINED_HEADERS or retained_f.max_row != 18 or retained_f.max_column != 5:
        result["retained_errors"].append(f"structure={headers(retained_f)}/{retained_f.max_row}x{retained_f.max_column}")
    for row_index, (year, revenue) in enumerate(expected_retained, start=2):
        got_values = [retained_v.cell(row_index, column).value for column in range(1, 6)]
        previous = None if row_index == 2 else expected_retained[row_index - 3][1]
        wanted = [
            year, revenue, revenue - THRESHOLD,
            None if previous is None else revenue - previous,
            None if previous is None else Decimal(revenue) / Decimal(previous) - Decimal("1"),
        ]
        for column, (got, expected) in enumerate(zip(got_values, wanted, strict=True), start=1):
            valid = got is None if expected is None else approx(got, expected, 1e-7)
            if not valid:
                result["retained_errors"].append(f"{retained_v.cell(row_index, column).coordinate}={got!r}")
        if retained_f.cell(row_index, 3).data_type != "f":
            result["formula_errors"].append(f"Retained_Years!C{row_index}:not_formula")
        if row_index > 2:
            for column in (4, 5):
                if retained_f.cell(row_index, column).data_type != "f":
                    result["formula_errors"].append(f"Retained_Years!{retained_f.cell(row_index, column).coordinate}:not_formula")

    removed_f = formulas["Removed_Years"]
    removed_v = values["Removed_Years"]
    expected_removed = truth["removed"]
    if headers(removed_f) != REMOVED_HEADERS or removed_f.max_row != 4 or removed_f.max_column != 3:
        result["removed_errors"].append(f"structure={headers(removed_f)}/{removed_f.max_row}x{removed_f.max_column}")
    for row_index, (year, revenue) in enumerate(expected_removed, start=2):
        wanted = [year, revenue, THRESHOLD - revenue]
        for column, expected in enumerate(wanted, start=1):
            got = removed_v.cell(row_index, column).value
            if not approx(got, expected, 1e-7):
                result["removed_errors"].append(f"{removed_v.cell(row_index, column).coordinate}={got!r}")
        if removed_f.cell(row_index, 3).data_type != "f":
            result["formula_errors"].append(f"Removed_Years!C{row_index}:not_formula")

    summary_f = formulas["Summary"]
    summary_v = values["Summary"]
    labels = [normalize(summary_v.cell(row, 1).value) for row in range(2, 14)]
    if labels != SUMMARY_LABELS:
        result["summary_errors"].append(f"labels={labels}")
    expected_summary = [
        20, 17, 3, Decimal("0.85"), truth["retained_total"], truth["retained_average"],
        truth["highest_year"], truth["highest_revenue"], truth["lowest_year"], truth["lowest_revenue"],
        truth["first_last_change"], truth["first_last_pct"],
    ]
    for row_index, expected in enumerate(expected_summary, start=2):
        if summary_f.cell(row_index, 2).data_type != "f":
            result["formula_errors"].append(f"Summary!B{row_index}:not_formula")
        got = summary_v.cell(row_index, 2).value
        if not approx(got, expected, 1e-7):
            result["summary_errors"].append(f"B{row_index}={got!r} expected={expected}")

    for sheet in values.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and re.match(r"^#(?:NAME|VALUE|REF|DIV/0|N/A)", cell.value, flags=re.IGNORECASE):
                    result["formula_errors"].append(f"{sheet.title}!{cell.coordinate}={cell.value}")
    result["native"] = {
        "tables": sum(len(sheet.tables) for sheet in formulas.worksheets),
        "charts": sum(len(sheet._charts) for sheet in formulas.worksheets),
        "freeze": all(sheet.freeze_panes == "A2" for sheet in formulas.worksheets),
        "filters": all(bool(sheet.auto_filter.ref) for sheet in formulas.worksheets),
        "conditional_formats": sum(len(sheet.conditional_formatting) for sheet in formulas.worksheets),
    }
    return result


def word_text(document: Document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def unsupported_claims(text: str) -> list[str]:
    patterns = {
        "unsupported_cause": r"(?:declined|fell|increased|rose|grew).{0,35}(?:because of|due to|driven by|caused by).{0,50}(?:inflation|market|product|management|competition|economy)",
        "profit_confusion": r"(?:revenue|income).{0,20}(?:is|equals|represents).{0,20}(?:profit|earnings)",
        "forecast": r"(?:will|guaranteed to|expected to)\s+(?:grow|rise|fall|decline|reach)",
    }
    return [name for name, pattern in patterns.items() if re.search(pattern, text, flags=re.IGNORECASE)]


def parse_word(path: Path, truth: dict[str, object]) -> dict[str, object]:
    result: dict[str, object] = {"readable": False, "structure_errors": [], "fact_errors": [], "boundary_errors": [], "unsupported": [], "tables": 0, "headings": []}
    try:
        document = Document(path)
    except Exception as exc:
        result["error"] = f"{type(exc).__name__}:{exc}"
        return result
    result["readable"] = True
    result["tables"] = len(document.tables)
    result["headings"] = [normalize(paragraph.text) for paragraph in document.paragraphs if str(paragraph.style.name).startswith("Heading")]
    text = normalize(word_text(document))
    if "Revenue Threshold Screening Report" not in text:
        result["structure_errors"].append("wrong_title")
    for heading in ("Screening Rule", "Executive Summary", "Removed Years", "Retained Revenue Findings", "Evidence Boundary"):
        if heading not in result["headings"]:
            result["structure_errors"].append(f"missing_heading:{heading}")
    for phrase in ("greater than or equal to $2,000,000", "strictly below $2,000,000", "equality boundary"):
        if phrase.lower() not in text.lower():
            result["structure_errors"].append(f"missing_rule:{phrase}")
    required_facts = [
        "20 annual observations", "17 years", "removes 3", "85.00%", money(truth["retained_total"]), money(truth["retained_average"]),
        f"{truth['highest_year']} — {money(truth['highest_revenue'])}", f"{truth['lowest_year']} — {money(truth['lowest_revenue'])}",
        money(truth["first_last_change"]), f"{Decimal(truth['first_last_pct']):.2%}", "$1,552,163.00",
    ]
    for token in required_facts:
        if token.lower() not in text.lower():
            result["fact_errors"].append(f"missing:{token}")
    for year, revenue in truth["removed"]:
        for token in (str(year), money(revenue), money(THRESHOLD - revenue)):
            if token.lower() not in text.lower():
                result["fact_errors"].append(f"removed:{year}:missing:{token}")
    for phrase in ("endpoint comparison", "not a claim of a continuous trend", "use only company_revenues.xlsx", "without assigning causes", "predicting future performance"):
        if phrase.lower() not in text.lower():
            result["boundary_errors"].append(f"missing:{phrase}")
    result["unsupported"] = unsupported_claims(text)
    return result


def render_office(artifact_dir: Path) -> dict[str, object]:
    result: dict[str, object] = {"xlsx_pages": 0, "docx_pages": 0, "errors": [], "logs": []}
    with tempfile.TemporaryDirectory(prefix="revenue-judge-render-") as tmp:
        root = Path(tmp)
        output = root / "output"
        profile = root / "profile"
        output.mkdir()
        profile.mkdir()
        for name in REQUIRED_FILES:
            source = artifact_dir / name
            if not source.is_file():
                result["errors"].append(f"missing:{name}")
                continue
            command = ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf", "--outdir", str(output), str(source)]
            try:
                proc = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300)
            except Exception as exc:
                result["errors"].append(f"{name}:{type(exc).__name__}:{exc}")
                continue
            result["logs"].append(proc.stdout[-1000:])
            pdf = output / f"{Path(name).stem}.pdf"
            if proc.returncode != 0 or not pdf.is_file():
                result["errors"].append(f"{name}:conversion_failed")
                continue
            info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=60)
            match = re.search(r"^Pages:\s+(\d+)", info.stdout, flags=re.MULTILINE)
            pages = int(match.group(1)) if match else 0
            extracted = subprocess.run(["pdftotext", str(pdf), "-"], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=60).stdout
            if name == XLSX_NAME:
                result["xlsx_pages"] = pages
                for token in ("2004", "2023", "Revenue_Rank", "Retained revenue by year", "First_to_Last_Change_Pct"):
                    if token not in extracted:
                        result["errors"].append(f"xlsx_render_missing:{token}")
            else:
                result["docx_pages"] = pages
                for token in ("Revenue Threshold Screening Report", "Removed Years", "Evidence Boundary", "1,552,163"):
                    if token not in extracted:
                        result["errors"].append(f"docx_render_missing:{token}")
            if re.search(r"#(?:NAME|VALUE|REF|DIV/0|N/A)", extracted, flags=re.IGNORECASE):
                result["errors"].append(f"{name}:formula_error_visible")
    if result["xlsx_pages"] != 5:
        result["errors"].append(f"xlsx_pages={result['xlsx_pages']}")
    if not 2 <= result["docx_pages"] <= 3:
        result["errors"].append(f"docx_pages={result['docx_pages']}")
    return result


def suspicious_content(artifact_dir: Path) -> list[str]:
    findings: list[str] = []
    for path in artifact_dir.iterdir() if artifact_dir.is_dir() else []:
        if not path.is_file():
            continue
        try:
            if path.suffix.lower() in {".md", ".txt"}:
                text = path.read_text(encoding="utf-8", errors="ignore").lower()
            elif path.suffix.lower() in {".xlsx", ".docx"}:
                with zipfile.ZipFile(path) as package:
                    text = " ".join(package.read(name).decode("utf-8", errors="ignore") for name in package.namelist() if name.endswith(".xml")).lower()
            else:
                text = ""
        except Exception:
            continue
        for marker in SUSPICIOUS_MARKERS:
            if marker in text:
                findings.append(f"{path.name}:{marker}")
    return sorted(set(findings))


def criterion(identifier: str, score: float, evidence: object) -> dict[str, object]:
    description, weight, kind, cap = SPECS[identifier]
    raw: dict[str, object] = {"description": description, "weight": weight, "type": kind, "passed": score >= 0.999}
    if cap is not None:
        raw["score_cap_on_failure"] = cap
    return {"criterion_id": identifier, "score": round(float(score), 6), "evidence": evidence, "method": "deterministic", "raw": raw}


def evaluate(artifact_dir: Path) -> dict[str, object]:
    rows, truth, source_errors = load_truth()
    delivered = sorted(path.name for path in artifact_dir.iterdir() if path.is_file()) if artifact_dir.is_dir() else []
    exact_files = delivered == sorted(REQUIRED_FILES)
    packages = {name: inspect_package(artifact_dir / name) for name in REQUIRED_FILES}
    package_ok = all(not info["broken"] and not info["missing_targets"] and not info["macros"] and not info["external"] for info in packages.values())
    workbook = parse_workbook(artifact_dir / XLSX_NAME, rows, truth) if (artifact_dir / XLSX_NAME).is_file() and not source_errors else {
        "readable": False, "source_errors": ["unavailable"], "retained_errors": [], "removed_errors": [], "formula_errors": [], "summary_errors": [], "native": {},
    }
    word = parse_word(artifact_dir / DOCX_NAME, truth) if (artifact_dir / DOCX_NAME).is_file() and not source_errors else {
        "readable": False, "structure_errors": ["unavailable"], "fact_errors": [], "boundary_errors": [], "unsupported": [], "tables": 0, "headings": [],
    }
    render = render_office(artifact_dir) if exact_files else {"xlsx_pages": 0, "docx_pages": 0, "errors": ["incomplete"]}
    suspicious = suspicious_content(artifact_dir)
    source_ok = workbook.get("readable") and not workbook.get("source_errors")
    retained_ok = source_ok and not workbook.get("retained_errors")
    removed_ok = source_ok and not workbook.get("removed_errors")
    formulas_ok = retained_ok and removed_ok and not workbook.get("formula_errors")
    summary_ok = formulas_ok and not workbook.get("summary_errors")
    native = workbook.get("native", {})
    native_ok = native.get("tables", 0) >= 5 and native.get("charts", 0) >= 1 and native.get("freeze") is True and native.get("filters") is True and native.get("conditional_formats", 0) >= 3
    word_structure_ok = word.get("readable") and word.get("tables", 0) >= 4 and not word.get("structure_errors") and not packages[DOCX_NAME].get("media")
    word_facts_ok = word_structure_ok and not word.get("fact_errors")
    word_boundary_ok = word_facts_ok and not word.get("boundary_errors") and not word.get("unsupported")
    cross_ok = summary_ok and word_facts_ok
    safe_ok = exact_files and package_ok and not suspicious
    render_ok = not render.get("errors")
    unsupported = word.get("unsupported", [])

    results = [
        criterion("R001", 1.0 if exact_files and not source_errors and package_ok else 0.0, {"source_errors": source_errors, "delivered": delivered, "packages": packages}),
        criterion("R002", 1.0 if source_ok else 0.0, workbook.get("source_errors")),
        criterion("R003", 1.0 if retained_ok else 0.0, workbook.get("retained_errors")),
        criterion("R004", 1.0 if removed_ok else 0.0, workbook.get("removed_errors")),
        criterion("R005", 1.0 if formulas_ok else 0.0, workbook.get("formula_errors")),
        criterion("R006", 1.0 if summary_ok else 0.0, workbook.get("summary_errors")),
        criterion("R007", 1.0 if native_ok else 0.0, native),
        criterion("R008", 1.0 if word_structure_ok else 0.0, {"tables": word.get("tables"), "headings": word.get("headings"), "errors": word.get("structure_errors")}),
        criterion("R009", 1.0 if word_facts_ok else 0.0, word.get("fact_errors")),
        criterion("R010", 1.0 if word_boundary_ok else 0.0, {"boundary_errors": word.get("boundary_errors"), "unsupported": unsupported}),
        criterion("R011", 1.0 if cross_ok else 0.0, {"workbook": summary_ok, "word": word_facts_ok}),
        criterion("R012", 1.0 if render_ok and native_ok else 0.0, {"render": render, "native": native}),
        criterion("R013", 1.0 if source_ok and retained_ok and removed_ok and formulas_ok and summary_ok else 0.0, {"source": source_ok, "retained": retained_ok, "removed": removed_ok, "formulas": formulas_ok, "summary": summary_ok}),
        criterion("R014", 1.0 if cross_ok else 0.0, {"workbook": summary_ok, "word": word_facts_ok}),
        criterion("R015", 1.0 if safe_ok and formulas_ok and native_ok else 0.0, {"safe": safe_ok, "formulas": formulas_ok, "native": native_ok, "suspicious": suspicious}),
        criterion("R016", 1.0 if not unsupported else 0.0, unsupported),
    ]
    positive_points = sum(SPECS[item["criterion_id"]][1] * float(item["score"]) for item in results if SPECS[item["criterion_id"]][1] > 0)
    penalty_points = sum(SPECS[item["criterion_id"]][1] for item in results if SPECS[item["criterion_id"]][1] < 0 and float(item["score"]) < 0.999)
    reward = max(0.0, min(1.0, (positive_points + penalty_points) / 100.0))
    hurdle_failures: list[str] = []
    for item in results:
        identifier = item["criterion_id"]
        _, _, kind, cap = SPECS[identifier]
        if kind == "hurdle" and float(item["score"]) < 0.999:
            hurdle_failures.append(identifier)
            reward = min(reward, float(cap))
    reward = round(reward, 6)
    return {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": [sha256(artifact_dir / name) for name in REQUIRED_FILES if (artifact_dir / name).is_file()],
        "criterion_results": results, "reward": reward, "passed": reward >= 0.8 and not hurdle_failures,
        "hurdle_failures": hurdle_failures,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    result = evaluate(args.artifact_dir)
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{result['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
