#!/usr/bin/env python3
"""收入门槛交付包的 v2 任务专属分层负例。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Callable

from docx import Document
from openpyxl import Workbook, load_workbook


ROOT=Path(__file__).resolve().parents[1];GOLDEN=ROOT/"solution"/"Golden";INPUTS=ROOT/"environment"/"input_files"
REQUIRED=("Revenue_Filtered_Analysis.xlsx","Revenue_Threshold_Report.docx")
Mutator=Callable[[Path],None]


def _load_judge():
    path=Path(__file__).with_name("judge.py");spec=importlib.util.spec_from_file_location("revenue_v2_judge_for_negatives",path)
    if spec is None or spec.loader is None:raise RuntimeError("无法载入 v2 Judge")
    module=importlib.util.module_from_spec(spec);sys.modules[spec.name]=module;spec.loader.exec_module(module);return module


JUDGE=_load_judge()


def _copy(target:Path)->None:
    target.mkdir(parents=True,exist_ok=True)
    for name in REQUIRED:shutil.copy2(GOLDEN/name,target/name)


def _excel(sheet:str,cell:str,value:object)->Mutator:
    def apply(root:Path)->None:
        path=root/REQUIRED[0];book=load_workbook(path);book[sheet][cell]=value;book.save(path);book.close()
    return apply


def _remove_chart(root:Path)->None:
    path=root/REQUIRED[0];book=load_workbook(path);book["Summary"]._charts=[];book.save(path);book.close()


def _remove_freeze(root:Path)->None:
    path=root/REQUIRED[0];book=load_workbook(path);book["Threshold_Audit"].freeze_panes=None;book.save(path);book.close()


def _remove_table(root:Path)->None:
    path=root/REQUIRED[0];book=load_workbook(path);del book["Retained_Years"].tables["RetainedRevenueTable"];book.save(path);book.close()


def _remove_conditional(root:Path)->None:
    path=root/REQUIRED[0];book=load_workbook(path);book["Threshold_Audit"].conditional_formatting._cf_rules.clear();book.save(path);book.close()


def _word_replace(old:str,new:str)->Mutator:
    def apply(root:Path)->None:
        path=root/REQUIRED[1];document=Document(path);locations=list(document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:locations.extend(cell.paragraphs)
        for paragraph in locations:
            if old in paragraph.text:
                paragraph.text=paragraph.text.replace(old,new,1);document.save(path);return
        raise AssertionError(f"Word 未找到：{old}")
    return apply


def _word_table(table_index:int,row_index:int,column_index:int,value:str)->Mutator:
    def apply(root:Path)->None:
        path=root/REQUIRED[1];document=Document(path);document.tables[table_index].rows[row_index].cells[column_index].text=value;document.save(path)
    return apply


def _word_remove_table(root:Path)->None:
    path=root/REQUIRED[1];document=Document(path);table=document.tables[2];table._element.getparent().remove(table._element);document.save(path)


def _word_add(text:str)->Mutator:
    def apply(root:Path)->None:
        path=root/REQUIRED[1];document=Document(path);document.add_paragraph(text);document.save(path)
    return apply


def _moderate_threshold_domain_conflict(root:Path)->None:
    path=root/REQUIRED[0];book=load_workbook(path)
    audit=book["Threshold_Audit"]
    for row in range(2,22):audit.cell(row,3).value=3_000_000;audit.cell(row,4).value='=IF(B%d>=C%d,"Retain","Remove")'%(row,row);audit.cell(row,5).value=f"=B{row}-C{row}"
    retained=book["Retained_Years"]
    for row in range(2,19):retained.cell(row,1).value=2000+row;retained.cell(row,2).value=3_000_000
    removed=book["Removed_Years"]
    for row in range(2,5):removed.cell(row,1).value=1990+row;removed.cell(row,2).value=1_000_000
    summary=book["Summary"]
    for row in range(2,14):summary.cell(row,2).value=0
    book.save(path);book.close()
    doc=root/REQUIRED[1];document=Document(doc)
    for table in document.tables:
        for row in table.rows[1:]:
            for cell in row.cells:cell.text="0"
    document.paragraphs[3].text="Years with revenue greater than or equal to $3,000,000 are retained; lower years are removed."
    document.paragraphs[5].text="The source contains 20 observations; 0 years are retained and 20 removed."
    document.paragraphs[8].text="All years are removed under the revised threshold."
    document.paragraphs[10].text="Retained revenue total is $0.00."
    document.paragraphs[11].text="No endpoint comparison is available."
    document.save(doc)


def _severe_valid_incomplete(root:Path)->None:
    book=Workbook();book.remove(book.active)
    headers={"Source_Data":["Year","Revenue_USD"],"Threshold_Audit":JUDGE.V1.AUDIT_HEADERS,"Retained_Years":JUDGE.V1.RETAINED_HEADERS,"Removed_Years":JUDGE.V1.REMOVED_HEADERS,"Summary":["Metric","Value"]}
    for name,row in headers.items():sheet=book.create_sheet(name);sheet.append(row)
    book.save(root/REQUIRED[0]);book.close()
    document=Document();document.add_paragraph("Revenue Threshold Screening Report",style="Title");document.add_heading("Screening Rule",level=1);document.add_paragraph("Analysis is incomplete.");document.save(root/REQUIRED[1])


def _input_only(root:Path)->None:
    for path in list(root.iterdir()):
        if path.is_file():path.unlink()
    for source in INPUTS.iterdir():
        if source.is_file():shutil.copy2(source,root/source.name)


PARTIALS:dict[str,tuple[Mutator,bool]]={
    "source_year_wrong":(_excel("Source_Data","A2",2003),False),
    "source_revenue_wrong":(_excel("Source_Data","B4",1_293_225),False),
    "audit_year_wrong":(_excel("Threshold_Audit","A7",2019),False),
    "audit_revenue_wrong":(_excel("Threshold_Audit","B15",2_463_500),False),
    "audit_threshold_wrong":(_excel("Threshold_Audit","C2",2_500_000),False),
    "audit_status_flattened":(_excel("Threshold_Audit","D4","Remove"),False),
    "audit_status_wrong_formula":(_excel("Threshold_Audit","D4",'=IF(B4>C4,"Retain","Remove")'),False),
    "audit_margin_flattened":(_excel("Threshold_Audit","E4",-806775),False),
    "audit_margin_wrong_reference":(_excel("Threshold_Audit","E4","=B4-C5"),False),
    "audit_rank_flattened":(_excel("Threshold_Audit","F2",3),False),
    "audit_rank_wrong_range":(_excel("Threshold_Audit","F2","=RANK(B2,$B$2:$B$20,0)"),False),
    "retained_year_wrong":(_excel("Retained_Years","A18",2024),False),
    "retained_revenue_wrong":(_excel("Retained_Years","B18",2_900_000),False),
    "retained_margin_flattened":(_excel("Retained_Years","C2",2_439_044),False),
    "retained_margin_wrong_formula":(_excel("Retained_Years","C2","=B2-2500000"),False),
    "retained_change_flattened":(_excel("Retained_Years","D3",-2_226_841),False),
    "retained_change_wrong_reference":(_excel("Retained_Years","D3","=B3-B1"),False),
    "retained_pct_flattened":(_excel("Retained_Years","E3",-0.501647),False),
    "retained_pct_wrong_formula":(_excel("Retained_Years","E3","=B3/B2"),False),
    "removed_year_wrong":(_excel("Removed_Years","A2",2005),False),
    "removed_revenue_wrong":(_excel("Removed_Years","B3",1_891_112),False),
    "removed_shortfall_flattened":(_excel("Removed_Years","C2",806775),False),
    "removed_shortfall_wrong_formula":(_excel("Removed_Years","C2","=2000000-B3"),False),
    "summary_label_wrong":(_excel("Summary","A2","Years_Total"),False),
    "summary_original_flattened":(_excel("Summary","B2",20),False),
    "summary_retained_wrong_formula":(_excel("Summary","B3","=16"),False),
    "summary_rate_flattened":(_excel("Summary","B5",.85),False),
    "summary_total_wrong_formula":(_excel("Summary","B6","=SUM(Retained_Years!B2:B17)"),False),
    "summary_average_flattened":(_excel("Summary","B7",3304272.65),False),
    "summary_high_year_wrong_formula":(_excel("Summary","B8","=2015"),False),
    "summary_low_revenue_wrong_formula":(_excel("Summary","B11","=2000000"),False),
    "summary_change_wrong_formula":(_excel("Summary","B12","=Retained_Years!B17-Retained_Years!B2"),False),
    "summary_change_pct_flattened":(_excel("Summary","B13",-.361902022147111),False),
    "native_chart_removed":(_remove_chart,False),
    "native_freeze_removed":(_remove_freeze,False),
    "native_table_removed":(_remove_table,False),
    "native_conditional_format_removed":(_remove_conditional,False),
    "word_heading_wrong":(_word_replace("Evidence Boundary","Evidence Notes"),False),
    "word_threshold_rule_wrong":(_word_replace("greater than or equal to $2,000,000","strictly greater than $2,000,000"),False),
    "word_retained_count_wrong":(_word_table(0,1,1,"16"),False),
    "word_retention_rate_wrong":(_word_table(0,1,3,"80.00%"),False),
    "word_removed_year_wrong":(_word_table(2,1,0,"2005"),False),
    "word_removed_shortfall_wrong":(_word_table(2,2,2,"$308,888.00"),False),
    "word_total_conflict":(_word_table(1,1,1,"$56,000,000.00"),False),
    "word_top_rank_wrong":(_word_table(3,1,1,"2015"),False),
    "word_endpoint_wrong":(_word_replace("-$1,606,499.00","-$1,500,000.00"),False),
    "word_boundary_missing":(_word_replace("without assigning causes or predicting future performance","without further comment"),False),
    "word_removed_table_missing":(_word_remove_table,False),
    "unsupported_causal_claim":(_word_add("Revenue fell because of inflation and market competition."),False),
    "unsupported_forecast":(_word_add("Revenue will grow next year."),False),
    "unexpected_debug_file":(lambda root:(root/"debug.txt").write_text("temporary\n",encoding="utf-8"),False),
    "moderate_threshold_domain_conflict":(_moderate_threshold_domain_conflict,True),
    "severe_valid_but_incomplete_package":(_severe_valid_incomplete,True),
}


def main()->None:
    parser=argparse.ArgumentParser();parser.add_argument("--report",type=Path);args=parser.parse_args();golden_recalc,_=JUDGE._recalculate(GOLDEN/REQUIRED[0]);golden_render=JUDGE.V1.render_office(GOLDEN);original_recalc,original_render=JUDGE._recalculate,JUDGE.V1.render_office
    assert JUDGE.evaluate(GOLDEN)["reward"]==1.0
    def fast_recalc(path:Path):
        book=load_workbook(path,data_only=False);values={}
        for sheet in book.worksheets:
            rows=[]
            for row_index,row in enumerate(sheet.iter_rows(),1):
                current=[]
                for column_index,cell in enumerate(row,1):
                    value=cell.value
                    if isinstance(value,str) and value.startswith("="):value=JUDGE._value(golden_recalc,sheet.title,row_index,column_index)
                    current.append(value)
                rows.append(current)
            values[sheet.title]=rows
        book.close();return values,"任务专属 Golden 公式缓存值"
    def evaluate_case(case:Path,real:bool)->float:
        if real:return float(JUDGE.evaluate(case)["reward"])
        JUDGE._recalculate=fast_recalc;JUDGE.V1.render_office=lambda _artifact:golden_render
        try:return float(JUDGE.evaluate(case)["reward"])
        finally:JUDGE._recalculate,JUDGE.V1.render_office=original_recalc,original_render
    records=[]
    with tempfile.TemporaryDirectory(prefix="revenue-v2-negatives-") as raw:
        base=Path(raw)
        for name in REQUIRED:
            for mode in ("missing","empty","corrupt"):
                case=base/f"fatal-{mode}-{name}";_copy(case);target=case/name
                if mode=="missing":target.unlink()
                elif mode=="empty":target.write_bytes(b"")
                else:target.write_bytes(b"not-a-valid-task-artifact")
                reward=evaluate_case(case,False);assert reward==0,(mode,name,reward);records.append({"case":f"{mode}:{name}","layer":"fatal_zero","reward":reward})
        empty=base/"fatal-empty";empty.mkdir();reward=evaluate_case(empty,False);assert reward==0;records.append({"case":"empty-output","layer":"fatal_zero","reward":reward})
        input_only=base/"fatal-input-only";_copy(input_only);_input_only(input_only);reward=evaluate_case(input_only,False);assert reward==0;records.append({"case":"unmodified-input-only","layer":"fatal_zero","reward":reward})
        for name,(mutator,real) in PARTIALS.items():
            case=base/f"partial-{name}";_copy(case);mutator(case);reward=evaluate_case(case,real);assert 0<reward<1,(name,reward);records.append({"case":name,"layer":"partial","reward":reward})
    partial=[float(record["reward"]) for record in records if record["layer"]=="partial"];by_name={str(record["case"]):float(record["reward"]) for record in records}
    assert by_name["moderate_threshold_domain_conflict"]<.8
    assert 0<by_name["severe_valid_but_incomplete_package"]<by_name["moderate_threshold_domain_conflict"]
    report={"total":len(records),"fatal_zero":sum(record["layer"]=="fatal_zero" for record in records),"partial_continuous":len(partial),"partial_reward_range":[min(partial),max(partial)],"moderate_threshold_domain_conflict":by_name["moderate_threshold_domain_conflict"],"severe_valid_but_incomplete":by_name["severe_valid_but_incomplete_package"],"cases":records}
    if args.report:args.report.parent.mkdir(parents=True,exist_ok=True);args.report.write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=="__main__":main()
