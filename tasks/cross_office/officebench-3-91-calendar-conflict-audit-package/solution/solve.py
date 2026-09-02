#!/usr/bin/env python3
"""根据逐条审计的 OfficeBench 3-91 日历真值构建 Golden。"""

from __future__ import annotations

import argparse
import hashlib
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


XLSX_NAME = "Calendar_Conflict_Audit.xlsx"
ICS_NAME = "Merged_Schedule.ics"
DOCX_NAME = "Calendar_Conflict_Report.docx"
SOURCE_HASHES = {
    "Bob.ics": "a641a03140fc0903fb3d3a320642c146edc5689ffae06712c84cb4b5ee60be98",
    "Tom.ics": "4f946c13f4ee53deb6d74b70f73da42cc9a1063fc1183b7daca4a44bf69eba99",
}
FIXED_DT = datetime(2026, 9, 1, 6, 30, 0, tzinfo=timezone.utc)
SEARCH_START = datetime(2024, 5, 1, 8, 0, 0, tzinfo=timezone.utc)
SEARCH_END = datetime(2024, 5, 1, 20, 0, 0, tzinfo=timezone.utc)

NAVY = "17324D"
BLUE = "276B8E"
TEAL = "168AAD"
RED = "B5403C"
WHITE = "FFFFFF"
PALE_BLUE = "EAF3F8"
PALE_GREEN = "E7F4EC"
PALE_AMBER = "FFF4D6"
PALE_RED = "FBE9E7"
LINE = "CAD6E0"
GRAY = "607080"


@dataclass(frozen=True)
class Event:
    owner: str
    summary: str
    start: datetime
    end: datetime
    uid: str
    source_file: str


# 以下 10 条记录来自对 Bob.ics 和 Tom.ics 的逐条人工核对。
EXPECTED_EVENTS = (
    ("Bob", "sleeping", "20240501T080000Z", "20240501T110000Z", "7cf33664-81fa-44da-860a-6076c22bbf70@7cf3.org", "Bob.ics"),
    ("Tom", "class", "20240501T080000Z", "20240501T115000Z", "fa6c3555-0e2e-4c56-b5b6-d2ec860444b7@fa6c.org", "Tom.ics"),
    ("Bob", "lunch", "20240501T120000Z", "20240501T130000Z", "d5a507d5-9f22-479f-88c2-5fe36727a71e@d5a5.org", "Bob.ics"),
    ("Tom", "lunch with David", "20240501T123000Z", "20240501T133000Z", "f2c7bb64-face-48ff-a494-c23a5b31b6ad@f2c7.org", "Tom.ics"),
    ("Bob", "nap", "20240501T130000Z", "20240501T140000Z", "7cad3249-9761-48fd-b1b8-6336a5136eee@7cad.org", "Bob.ics"),
    ("Tom", "report to Charlie", "20240501T140000Z", "20240501T143000Z", "33d7143a-de13-4041-aa20-a88866b3acad@33d7.org", "Tom.ics"),
    ("Bob", "class", "20240501T160000Z", "20240501T180000Z", "6f8db951-e0c8-47f5-bb10-72a3491734ba@6f8d.org", "Bob.ics"),
    ("Tom", "class", "20240501T160000Z", "20240501T180000Z", "2fc598fd-361b-46c0-8b12-426c4d7939af@2fc5.org", "Tom.ics"),
    ("Bob", "dinner", "20240501T183000Z", "20240501T193000Z", "f825c769-a892-426d-8aff-31e2801edfe2@f825.org", "Bob.ics"),
    ("Tom", "dinner", "20240501T183000Z", "20240501T193000Z", "13e17887-2bbc-42a2-a881-5aabcdfda201@13e1.org", "Tom.ics"),
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def parse_utc(value: str) -> datetime:
    return datetime.strptime(value, "%Y%m%dT%H%M%SZ").replace(tzinfo=timezone.utc)


def format_utc(value: datetime) -> str:
    return value.astimezone(timezone.utc).strftime("%Y%m%dT%H%M%SZ")


def excel_datetime(value: datetime) -> str:
    """用原生 DATE+TIME 公式写入 UTC 时刻。

    LibreOffice 二次保存 Python datetime 时，对 11:50 这类二进制小数可能显示为
    11:49:59。明示 DATE/TIME 组件可保留冻结来源的整秒语义。
    """
    utc = value.astimezone(timezone.utc)
    return f"=DATE({utc.year},{utc.month},{utc.day})+TIME({utc.hour},{utc.minute},{utc.second})"


def unfold_ics(text: str) -> list[str]:
    lines: list[str] = []
    for raw in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        if raw.startswith((" ", "\t")) and lines:
            lines[-1] += raw[1:]
        elif raw:
            lines.append(raw)
    return lines


def parse_calendar(path: Path, owner: str) -> list[Event]:
    events: list[Event] = []
    current: dict[str, str] | None = None
    for line in unfold_ics(path.read_text(encoding="utf-8")):
        if line == "BEGIN:VEVENT":
            current = {}
            continue
        if line == "END:VEVENT":
            if current is None:
                raise ValueError(f"{path.name} VEVENT 边界异常")
            required = {"DTSTART", "DTEND", "SUMMARY", "UID"}
            if not required.issubset(current):
                raise ValueError(f"{path.name} VEVENT 缺字段：{required - set(current)}")
            events.append(Event(owner, current["SUMMARY"], parse_utc(current["DTSTART"]), parse_utc(current["DTEND"]), current["UID"], path.name))
            current = None
            continue
        if current is not None and ":" in line:
            key, value = line.split(":", 1)
            current[key.split(";", 1)[0].upper()] = value
    return events


def validate_and_load(input_dir: Path) -> list[Event]:
    events: list[Event] = []
    for filename, owner in (("Bob.ics", "Bob"), ("Tom.ics", "Tom")):
        path = input_dir / filename
        if not path.is_file():
            raise FileNotFoundError(path)
        actual = sha256(path)
        if actual != SOURCE_HASHES[filename]:
            raise ValueError(f"{filename} SHA256 不一致：{actual}")
        events.extend(parse_calendar(path, owner))
    events.sort(key=lambda event: (event.start, event.owner, event.summary, event.uid))
    actual_truth = tuple((event.owner, event.summary, format_utc(event.start), format_utc(event.end), event.uid, event.source_file) for event in events)
    if actual_truth != EXPECTED_EVENTS:
        raise ValueError("冻结 ICS 事件与逐条人工真值不一致")
    if any(event.end <= event.start for event in events):
        raise ValueError("存在非正时长事件")
    return events


def overlap(left: Event, right: Event) -> tuple[datetime, datetime] | None:
    start = max(left.start, right.start)
    end = min(left.end, right.end)
    return (start, end) if start < end else None


def analyze(events: list[Event]) -> dict[str, object]:
    ids = {event: f"E{index:02d}" for index, event in enumerate(events, start=1)}
    bob = [event for event in events if event.owner == "Bob"]
    tom = [event for event in events if event.owner == "Tom"]
    internal = [(left, right, overlap(left, right)) for owner in (bob, tom) for index, left in enumerate(owner) for right in owner[index + 1 :] if overlap(left, right)]
    pairs = []
    for left in bob:
        for right in tom:
            interval = overlap(left, right)
            if interval:
                pairs.append((left, right, interval[0], interval[1]))
    pairs.sort(key=lambda row: (row[2], row[0].summary, row[1].summary))

    conflict_windows: list[tuple[datetime, datetime]] = []
    for _, _, start, end in pairs:
        if not conflict_windows or start > conflict_windows[-1][1]:
            conflict_windows.append((start, end))
        else:
            prior_start, prior_end = conflict_windows[-1]
            conflict_windows[-1] = (prior_start, max(prior_end, end))

    busy: list[tuple[datetime, datetime]] = []
    for event in events:
        start = max(event.start, SEARCH_START)
        end = min(event.end, SEARCH_END)
        if start >= end:
            continue
        if not busy or start > busy[-1][1]:
            busy.append((start, end))
        else:
            busy[-1] = (busy[-1][0], max(busy[-1][1], end))
    free: list[tuple[datetime, datetime]] = []
    cursor = SEARCH_START
    for start, end in busy:
        if cursor < start:
            free.append((cursor, start))
        cursor = max(cursor, end)
    if cursor < SEARCH_END:
        free.append((cursor, SEARCH_END))

    if internal or len(pairs) != 5 or len(conflict_windows) != 4 or len(free) != 4:
        raise ValueError("人工审计真值自检失败")
    conflict_minutes = sum(int((end - start).total_seconds() // 60) for start, end in conflict_windows)
    free_minutes = sum(int((end - start).total_seconds() // 60) for start, end in free)
    if conflict_minutes != 420 or free_minutes != 160 or max(int((end - start).total_seconds() // 60) for start, end in free) != 90:
        raise ValueError("冲突或空闲窗口真值自检失败")
    return {"ids": ids, "internal": internal, "pairs": pairs, "conflict_windows": conflict_windows, "free": free}


def style_sheet(sheet, widths: dict[str, float]) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 34
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color=LINE)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=9, color=NAVY)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:1"
    sheet.page_margins.left = 0.2
    sheet.page_margins.right = 0.2
    sheet.page_margins.top = 0.3
    sheet.page_margins.bottom = 0.3


def add_table(sheet, name: str, ref: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)


def recalculate_workbook(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="calendar-xlsx-recalc-") as tmp:
        root = Path(tmp)
        source_dir = root / "source"
        output_dir = root / "output"
        profile = root / "profile"
        source_dir.mkdir()
        output_dir.mkdir()
        profile.mkdir()
        source = source_dir / path.name
        shutil.copy2(path, source)
        proc = subprocess.run(
            ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "xlsx", "--outdir", str(output_dir), str(source)],
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            timeout=300,
        )
        converted = output_dir / path.name
        if proc.returncode != 0 or not converted.is_file():
            raise RuntimeError(f"LibreOffice 公式重算失败：{proc.stdout[-1000:]}")
        shutil.copy2(converted, path)


def restore_frozen_rows(path: Path, sheet_count: int) -> None:
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with tempfile.TemporaryDirectory(prefix="calendar-xlsx-freeze-") as tmp:
        target = Path(tmp) / path.name
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w") as destination:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename in {f"xl/worksheets/sheet{index}.xml" for index in range(1, sheet_count + 1)}:
                    root = etree.fromstring(data)
                    sheet_view = root.find(f"{{{namespace}}}sheetViews/{{{namespace}}}sheetView")
                    if sheet_view is not None:
                        for existing in sheet_view.findall(f"{{{namespace}}}pane"):
                            sheet_view.remove(existing)
                        pane = etree.Element(f"{{{namespace}}}pane")
                        pane.set("ySplit", "1")
                        pane.set("topLeftCell", "A2")
                        pane.set("activePane", "bottomLeft")
                        pane.set("state", "frozen")
                        sheet_view.insert(0, pane)
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                destination.writestr(info, data)
        shutil.copy2(target, path)


def build_workbook(events: list[Event], analysis: dict[str, object], output_path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.properties.creator = "EvoCowork OfficeBench adapter"
    workbook.properties.title = "Calendar Conflict Audit"
    workbook.properties.created = FIXED_DT.replace(tzinfo=None)
    workbook.properties.modified = FIXED_DT.replace(tzinfo=None)
    ids: dict[Event, str] = analysis["ids"]  # type: ignore[assignment]
    pairs: list[tuple[Event, Event, datetime, datetime]] = analysis["pairs"]  # type: ignore[assignment]
    conflict_windows: list[tuple[datetime, datetime]] = analysis["conflict_windows"]  # type: ignore[assignment]
    free_windows: list[tuple[datetime, datetime]] = analysis["free"]  # type: ignore[assignment]

    source = workbook.create_sheet("Source_Events")
    source.append(["Event_ID", "Owner", "Summary", "DTSTART_UTC", "DTEND_UTC", "Duration_Minutes", "UID", "Source_File", "Internal_Conflict_Count", "Cross_Owner_Conflict_Count"])
    for row_index, event in enumerate(events, start=2):
        source.append([
            ids[event], event.owner, event.summary, excel_datetime(event.start), excel_datetime(event.end),
            f"=ROUND((E{row_index}-D{row_index})*1440,0)", event.uid, event.source_file,
            f'=SUMPRODUCT(($B$2:$B$11=B{row_index})*($A$2:$A$11<>A{row_index})*($D$2:$D$11<E{row_index})*($E$2:$E$11>D{row_index}))',
            f'=SUMPRODUCT(($B$2:$B$11<>B{row_index})*($D$2:$D$11<E{row_index})*($E$2:$E$11>D{row_index}))',
        ])
    style_sheet(source, {"A": 9, "B": 9, "C": 21, "D": 20, "E": 20, "F": 17, "G": 30, "H": 13, "I": 19, "J": 22})
    for row in range(2, 12):
        source.row_dimensions[row].height = 25
        # 列名已明确标注 UTC，打印版使用分钟精度以避免日期值因列宽不足显示为 ###。
        source.cell(row, 4).number_format = "yyyy-mm-dd hh:mm"
        source.cell(row, 5).number_format = "yyyy-mm-dd hh:mm"
        source.cell(row, 6).number_format = "0"
    source.conditional_formatting.add("I2:I11", FormulaRule(formula=["I2>0"], fill=PatternFill("solid", fgColor=PALE_RED)))
    source.conditional_formatting.add("J2:J11", FormulaRule(formula=["J2>0"], fill=PatternFill("solid", fgColor=PALE_AMBER)))
    add_table(source, "CalendarSourceEventsTable", "A1:J11")
    source.print_area = "A1:J11"

    pair_sheet = workbook.create_sheet("Pairwise_Conflicts")
    pair_sheet.append(["Conflict_ID", "Bob_Event_ID", "Bob_Summary", "Tom_Event_ID", "Tom_Summary", "Overlap_Start_UTC", "Overlap_End_UTC", "Overlap_Minutes", "Same_Title"])
    for row_index, (bob, tom, start, end) in enumerate(pairs, start=2):
        pair_sheet.append([
            f"P{row_index - 1:02d}", ids[bob], bob.summary, ids[tom], tom.summary,
            excel_datetime(start), excel_datetime(end), f"=ROUND((G{row_index}-F{row_index})*1440,0)", f'=IF(C{row_index}=E{row_index},"Yes","No")',
        ])
    style_sheet(pair_sheet, {"A": 12, "B": 14, "C": 22, "D": 14, "E": 24, "F": 23, "G": 23, "H": 18, "I": 14})
    for row in range(2, 7):
        pair_sheet.cell(row, 6).number_format = "yyyy-mm-dd hh:mm"
        pair_sheet.cell(row, 7).number_format = "yyyy-mm-dd hh:mm"
        pair_sheet.cell(row, 8).number_format = "0"
    pair_sheet.conditional_formatting.add("I2:I6", FormulaRule(formula=['I2="Yes"'], fill=PatternFill("solid", fgColor=PALE_BLUE)))
    add_table(pair_sheet, "CalendarPairwiseConflictsTable", "A1:I6")
    pair_sheet.print_area = "A1:I6"

    window_sheet = workbook.create_sheet("Conflict_Windows")
    window_sheet.append(["Window_ID", "Start_UTC", "End_UTC", "Duration_Minutes", "Active_Bob_Events", "Active_Tom_Events", "Pairwise_Conflict_Count"])
    for row_index, (start, end) in enumerate(conflict_windows, start=2):
        bob_names = sorted({event.summary for event in events if event.owner == "Bob" and event.start < end and event.end > start})
        tom_names = sorted({event.summary for event in events if event.owner == "Tom" and event.start < end and event.end > start})
        window_sheet.append([
            f"W{row_index - 1:02d}", excel_datetime(start), excel_datetime(end), f"=ROUND((C{row_index}-B{row_index})*1440,0)",
            "; ".join(bob_names), "; ".join(tom_names),
            f'=COUNTIFS(Pairwise_Conflicts!$F$2:$F$6,"<"&C{row_index},Pairwise_Conflicts!$G$2:$G$6,">"&B{row_index})',
        ])
    style_sheet(window_sheet, {"A": 12, "B": 23, "C": 23, "D": 18, "E": 32, "F": 32, "G": 25})
    for row in range(2, 6):
        window_sheet.cell(row, 2).number_format = "yyyy-mm-dd hh:mm"
        window_sheet.cell(row, 3).number_format = "yyyy-mm-dd hh:mm"
        window_sheet.cell(row, 4).number_format = "0"
    window_sheet.conditional_formatting.add("D2:D5", FormulaRule(formula=["D2>=120"], fill=PatternFill("solid", fgColor=PALE_RED)))
    add_table(window_sheet, "CalendarConflictWindowsTable", "A1:G5")
    window_sheet.print_area = "A1:G5"

    free_sheet = workbook.create_sheet("Joint_Free_Windows")
    free_sheet.append(["Window_ID", "Search_Window_Date", "Start_UTC", "End_UTC", "Duration_Minutes", "Is_Longest", "Recommendation"])
    for row_index, (start, end) in enumerate(free_windows, start=2):
        free_sheet.append([
            f"F{row_index - 1:02d}", "2024-05-01", excel_datetime(start), excel_datetime(end), f"=ROUND((D{row_index}-C{row_index})*1440,0)",
            f'=IF(E{row_index}=MAX($E$2:$E$5),"Yes","No")',
            f'=IF(F{row_index}="Yes","Recommended longest joint window","Available alternative")',
        ])
    style_sheet(free_sheet, {"A": 12, "B": 20, "C": 23, "D": 23, "E": 18, "F": 14, "G": 35})
    for row in range(2, 6):
        free_sheet.cell(row, 3).number_format = "yyyy-mm-dd hh:mm"
        free_sheet.cell(row, 4).number_format = "yyyy-mm-dd hh:mm"
        free_sheet.cell(row, 5).number_format = "0"
    free_sheet.conditional_formatting.add("F2:F5", FormulaRule(formula=['F2="Yes"'], fill=PatternFill("solid", fgColor=PALE_GREEN)))
    add_table(free_sheet, "CalendarJointFreeWindowsTable", "A1:G5")
    free_sheet.print_area = "A1:G5"

    summary = workbook.create_sheet("Summary")
    summary.append(["Metric", "Value", "Evidence"])
    metrics = [
        ("Total_Events", "=COUNTA(Source_Events!A2:A11)", "Two frozen calendars"),
        ("Bob_Events", '=COUNTIF(Source_Events!B2:B11,"Bob")', "Source_Events owner count"),
        ("Tom_Events", '=COUNTIF(Source_Events!B2:B11,"Tom")', "Source_Events owner count"),
        ("Internal_Conflict_Pairs", "=SUM(Source_Events!I2:I11)/2", "Half-open intervals within each owner"),
        ("Cross_Owner_Pairwise_Conflicts", "=COUNTA(Pairwise_Conflicts!A2:A6)", "Strict positive overlaps"),
        ("Merged_Conflict_Windows", "=COUNTA(Conflict_Windows!A2:A5)", "Overlapping or touching pairwise intervals merged"),
        ("Total_Conflict_Window_Minutes", "=SUM(Conflict_Windows!D2:D5)", "Union duration; avoids double counting"),
        ("Joint_Free_Windows", "=COUNTA(Joint_Free_Windows!A2:A5)", "Within 08:00-20:00 UTC"),
        ("Total_Joint_Free_Minutes", "=SUM(Joint_Free_Windows!E2:E5)", "Complement of combined busy intervals"),
        ("Longest_Joint_Free_Minutes", "=MAX(Joint_Free_Windows!E2:E5)", "Recommended window"),
        ("Same_Title_Same_Time_Pairs", '=COUNTIF(Pairwise_Conflicts!I2:I6,"Yes")', "Class and dinner"),
        ("Source_Files", '=SUMPRODUCT(1/COUNTIF(Source_Events!H2:H11,Source_Events!H2:H11))', "Bob.ics and Tom.ics"),
        ("Merged_ICS_Events", "=COUNTA(Source_Events!A2:A11)", "All source events retained"),
    ]
    for row in metrics:
        summary.append(row)
    summary["D1"] = "Category"
    summary["E1"] = "Minutes"
    summary["D2"] = "Conflict windows"
    summary["E2"] = "=B8"
    summary["D3"] = "Joint free windows"
    summary["E3"] = "=B10"
    style_sheet(summary, {"A": 38, "B": 20, "C": 52, "D": 22, "E": 16, "F": 3, "G": 16, "H": 16, "I": 16, "J": 16})
    for row in range(2, 15):
        summary.cell(row, 1).fill = PatternFill("solid", fgColor=PALE_BLUE)
        summary.cell(row, 1).font = Font(name="Aptos", size=9, bold=True, color=NAVY)
        summary.cell(row, 2).font = Font(name="Aptos Display", size=11, bold=True, color=TEAL)
        summary.cell(row, 2).number_format = "0"
    summary.conditional_formatting.add("B5:B8", FormulaRule(formula=["B5>0"], fill=PatternFill("solid", fgColor=PALE_AMBER)))
    add_table(summary, "CalendarAuditSummaryTable", "A1:C14")
    chart = BarChart()
    chart.style = 10
    chart.title = "Conflict vs joint free minutes"
    chart.y_axis.title = "Minutes"
    chart.x_axis.title = "Window type"
    chart.height = 7.2
    chart.width = 13.5
    chart.add_data(Reference(summary, min_col=5, min_row=1, max_row=3), titles_from_data=True)
    chart.series[0].tx = SeriesLabel(v="Minutes")
    chart.set_categories(Reference(summary, min_col=4, min_row=2, max_row=3))
    summary.add_chart(chart, "G2")
    summary.print_area = "A1:O14"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    recalculate_workbook(output_path)
    restore_frozen_rows(output_path, 5)
    normalize_ooxml(output_path)


def build_merged_ics(events: list[Event], output_path: Path) -> None:
    lines = [
        "BEGIN:VCALENDAR",
        "VERSION:2.0",
        "PRODID:-//EvoCowork//OfficeBench 3-91 Calendar Audit//EN",
        "CALSCALE:GREGORIAN",
        "METHOD:PUBLISH",
        "X-WR-CALNAME:Bob and Tom Merged Schedule",
    ]
    for event in events:
        lines.extend([
            "BEGIN:VEVENT",
            f"UID:{event.uid}",
            f"DTSTAMP:{format_utc(FIXED_DT)}",
            f"DTSTART:{format_utc(event.start)}",
            f"DTEND:{format_utc(event.end)}",
            f"SUMMARY:{event.summary}",
            f"X-EVOCOWORK-OWNER:{event.owner}",
            f"X-EVOCOWORK-SOURCE:{event.source_file}",
            "END:VEVENT",
        ])
    lines.append("END:VCALENDAR")
    output_path.write_bytes(("\r\n".join(lines) + "\r\n").encode("utf-8"))


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE, size: str = "6") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_width(cell, centimeters: float) -> None:
    properties = cell._tc.get_or_add_tcPr()
    width = properties.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        properties.append(width)
    width.set(qn("w:w"), str(int(centimeters * 567)))
    width.set(qn("w:type"), "dxa")


def add_word_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float], font_size: float = 7.4) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    for index, value in enumerate(headers):
        cell = header.cells[index]
        cell.text = value
        set_cell_width(cell, widths[index])
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.name = "Aptos"
            run.font.size = Pt(font_size)
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(WHITE)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.text = value
            set_cell_width(cell, widths[index])
            set_cell_border(cell)
            if row_index % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = RGBColor.from_string(NAVY)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def build_report(events: list[Event], analysis: dict[str, object], output_path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.45)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9)
    styles["Normal"].font.color.rgb = RGBColor.from_string(NAVY)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    for level, size in ((1, 15), (2, 12)):
        style = styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE if level == 1 else TEAL)
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(4)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Calendar Conflict Audit Report")
    run.font.name = "Aptos Display"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = document.add_paragraph("Bob.ics + Tom.ics | 2024-05-01 UTC | Deterministic interval audit")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.color.rgb = RGBColor.from_string(GRAY)

    document.add_heading("Scope and Time Semantics", level=1)
    document.add_paragraph(
        "This audit reconciles two frozen iCalendar files containing 10 VEVENT records, five for Bob and five for Tom. "
        "Every DTSTART and DTEND ends in Z and is therefore interpreted as UTC without local-time conversion. "
        "Conflicts use half-open intervals [DTSTART, DTEND): touching endpoints are not conflicts."
    )

    document.add_heading("Source Event Reconciliation", level=1)
    event_rows = []
    ids: dict[Event, str] = analysis["ids"]  # type: ignore[assignment]
    for event in events:
        event_rows.append([
            ids[event], event.owner, event.summary,
            event.start.strftime("%H:%M"), event.end.strftime("%H:%M"),
            str(int((event.end - event.start).total_seconds() // 60)), event.source_file,
        ])
    add_word_table(document, ["ID", "Owner", "Summary", "Start", "End", "Min", "Source"], event_rows, [1.4, 1.7, 4.5, 2.0, 2.0, 1.4, 2.6], font_size=6.8)

    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_heading("Conflict Findings", level=1)
    document.add_paragraph(
        "Bob has 0 internal conflict pairs and Tom has 0 internal conflict pairs. Across owners, five event pairs overlap. "
        "Their union is four merged conflict windows totaling 420 minutes; this union total avoids double-counting the two adjacent lunch/nap pairwise overlaps."
    )
    pairs: list[tuple[Event, Event, datetime, datetime]] = analysis["pairs"]  # type: ignore[assignment]
    pair_rows = []
    for index, (bob, tom, start, end) in enumerate(pairs, start=1):
        pair_rows.append([
            f"P{index:02d}", f"{ids[bob]} {bob.summary}", f"{ids[tom]} {tom.summary}",
            start.strftime("%H:%M"), end.strftime("%H:%M"), str(int((end - start).total_seconds() // 60)),
            "Yes" if bob.summary == tom.summary else "No",
        ])
    add_word_table(document, ["Pair", "Bob event", "Tom event", "Start", "End", "Min", "Same title"], pair_rows, [1.4, 4.0, 4.2, 1.8, 1.8, 1.3, 2.0], font_size=7.0)

    conflict_windows: list[tuple[datetime, datetime]] = analysis["conflict_windows"]  # type: ignore[assignment]
    window_rows = []
    for index, (start, end) in enumerate(conflict_windows, start=1):
        bob_names = "; ".join(sorted({event.summary for event in events if event.owner == "Bob" and event.start < end and event.end > start}))
        tom_names = "; ".join(sorted({event.summary for event in events if event.owner == "Tom" and event.start < end and event.end > start}))
        count = sum(1 for _, _, pair_start, pair_end in pairs if pair_start < end and pair_end > start)
        window_rows.append([f"W{index:02d}", start.strftime("%H:%M"), end.strftime("%H:%M"), str(int((end - start).total_seconds() // 60)), bob_names, tom_names, str(count)])
    add_word_table(document, ["Window", "Start", "End", "Min", "Bob active", "Tom active", "Pairs"], window_rows, [1.5, 1.8, 1.8, 1.3, 4.0, 4.1, 1.4], font_size=6.9)

    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_heading("Joint Availability", level=1)
    document.add_paragraph(
        "Within the frozen 08:00-20:00 UTC search window, both calendars are simultaneously free in four windows totaling 160 minutes. "
        "The longest is 14:30-16:00 UTC (90 minutes) and is the evidence-based recommendation; no new meeting is created."
    )
    free_windows: list[tuple[datetime, datetime]] = analysis["free"]  # type: ignore[assignment]
    free_rows = []
    for index, (start, end) in enumerate(free_windows, start=1):
        minutes = int((end - start).total_seconds() // 60)
        free_rows.append([f"F{index:02d}", "2024-05-01", start.strftime("%H:%M"), end.strftime("%H:%M"), str(minutes), "Yes" if minutes == 90 else "No"])
    add_word_table(document, ["Window", "Date", "Start UTC", "End UTC", "Minutes", "Longest"], free_rows, [2.0, 3.0, 2.7, 2.7, 2.3, 2.3], font_size=7.4)

    document.add_heading("Merged Calendar Control", level=1)
    document.add_paragraph(
        "Merged_Schedule.ics retains exactly 10 VEVENT records. Every source UID, DTSTART, DTEND and SUMMARY is unchanged, "
        "and each event carries X-EVOCOWORK-OWNER and X-EVOCOWORK-SOURCE for traceability. Two pairs have the same title and the same time: class at 16:00-18:00 UTC and dinner at 18:30-19:30 UTC."
    )
    add_word_table(
        document,
        ["Control", "Expected", "Result"],
        [
            ["Source files", "Bob.ics + Tom.ics", "2 reconciled"],
            ["VEVENT count", "5 + 5", "10 retained"],
            ["UID/time/title changes", "None", "None"],
            ["Owner/source extensions", "Present on every event", "10/10"],
        ],
        [5.0, 6.8, 5.5],
        font_size=7.4,
    )

    document.add_heading("Audit Conclusion", level=1)
    document.add_paragraph(
        "Neither individual calendar is internally conflicting. The cross-owner audit finds five overlapping event pairs grouped into four conflict windows. "
        "The merged calendar is a faithful consolidation, not a rescheduling action. The 14:30-16:00 UTC window is the longest shared availability inside the defined search horizon. "
        "No timezone, attendee, location, priority or additional meeting is inferred."
    )

    add_page_number(document.sections[0].footer.paragraphs[0])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    normalize_ooxml(output_path)


def normalize_ooxml(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="calendar-ooxml-normalize-") as tmp:
        target = Path(tmp) / path.name
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(target, "w") as destination:
            for info in sorted(source.infolist(), key=lambda item: item.filename):
                data = source.read(info.filename)
                if info.filename == "docProps/core.xml":
                    root = etree.fromstring(data)
                    dcterms = "http://purl.org/dc/terms/"
                    for field in ("created", "modified"):
                        node = root.find(f"{{{dcterms}}}{field}")
                        if node is not None:
                            node.text = "2026-09-01T06:30:00Z"
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                elif info.filename.startswith("xl/charts/") and info.filename.endswith(".xml"):
                    root = etree.fromstring(data)
                    chart_ns = "http://schemas.openxmlformats.org/drawingml/2006/chart"
                    category_axis = root.find(f".//{{{chart_ns}}}catAx/{{{chart_ns}}}axId")
                    value_axis = root.find(f".//{{{chart_ns}}}valAx/{{{chart_ns}}}axId")
                    if category_axis is not None and value_axis is not None:
                        mapping = {category_axis.get("val", ""): "39117001", value_axis.get("val", ""): "39117002"}
                        for node in root.iter():
                            if etree.QName(node).localname in {"axId", "crossAx"} and node.get("val") in mapping:
                                node.set("val", mapping[node.get("val")])
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                copied = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
                copied.compress_type = zipfile.ZIP_DEFLATED
                copied.external_attr = info.external_attr
                copied.create_system = info.create_system
                copied.flag_bits = info.flag_bits
                destination.writestr(copied, data)
        shutil.copy2(target, path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    events = validate_and_load(args.input_dir)
    analysis = analyze(events)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for child in args.output_dir.iterdir():
        if child.is_file():
            child.unlink()
        elif child.is_dir():
            shutil.rmtree(child)
    build_workbook(events, analysis, args.output_dir / XLSX_NAME)
    build_merged_ics(events, args.output_dir / ICS_NAME)
    build_report(events, analysis, args.output_dir / DOCX_NAME)


if __name__ == "__main__":
    main()
