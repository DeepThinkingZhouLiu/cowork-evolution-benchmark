#!/usr/bin/env python3
"""OfficeBench 3-91 日历冲突审计包的 v2 任务专属原子评分。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from openpyxl import load_workbook


TASK_ID="cowork-evo/officebench-3-91-calendar-conflict-audit-package"
RUBRIC_VERSION="officebench-3-91-calendar-conflict-audit-v0.2";JUDGE_VERSION="calendar-conflict-audit-atomic-v2.0"
XLSX="Calendar_Conflict_Audit.xlsx";ICS="Merged_Schedule.ics";DOCX="Calendar_Conflict_Report.docx";REQUIRED=(XLSX,ICS,DOCX)
WEIGHTS={"R001":5,"R002":12,"R003":8,"R004":5,"R005":12,"R006":8,"R007":10,"R008":8,"R009":5,"R010":8,"R011":5,"R012":6,"R013":5,"R014":3}


def _load_v1():
    path=Path(__file__).with_name("judge_v1.py");spec=importlib.util.spec_from_file_location("calendar_judge_v1",path)
    if spec is None or spec.loader is None:raise RuntimeError("无法载入 v1 Judge")
    module=importlib.util.module_from_spec(spec);sys.modules[spec.name]=module;spec.loader.exec_module(module);return module


V1=_load_v1()


def _root()->Path:return Path(__file__).resolve().parents[1]


TRUTH,SOURCE_ERRORS=V1.truth_from_sources(_root())


def _sha256(path:Path)->str:
    digest=hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda:handle.read(1024*1024),b""):digest.update(block)
    return digest.hexdigest()


def _norm(value:Any)->str:return re.sub(r"\s+"," ","" if value is None else str(value)).strip().casefold()


def _same(actual:Any,expected:Any,tolerance:float=1e-6)->bool:
    if isinstance(expected,(int,float)) and not isinstance(expected,bool):
        try:return abs(float(actual)-float(expected))<=tolerance
        except (TypeError,ValueError):return False
    return _norm(actual)==_norm(expected)


def _time(actual:Any,expected:datetime)->bool:
    parsed=V1.as_utc(actual)
    return parsed is not None and abs((parsed-expected).total_seconds())<=.5


def _formula(value:Any,*tokens:str)->bool:
    if not isinstance(value,str) or not value.startswith("="):return False
    normalized=re.sub(r"[\s'$]","",value).upper()
    return all(re.sub(r"[\s'$]","",token).upper() in normalized for token in tokens)


def _ratio(items:Iterable[tuple[bool,str]])->tuple[float,list[str],int,int]:
    checks=list(items);earned=sum(bool(ok) for ok,_ in checks)
    return earned/max(1,len(checks)),[label for ok,label in checks if not ok],earned,len(checks)


def _criterion(identifier:str,items:Iterable[tuple[bool,str]])->dict[str,Any]:
    score,failures,earned,total=_ratio(items)
    return {"criterion_id":identifier,"score":round(score,6),"evidence":"全部任务专属原子检查通过" if not failures else "未通过："+"；".join(failures[:28]),"method":"deterministic","raw":{"earned_units":earned,"total_units":total,"failure_count":len(failures)}}


def _core(artifact:Path)->tuple[bool,list[str],list[str]]:
    errors=[];hashes=[]
    if not artifact.is_dir():return False,["输出目录不存在"],hashes
    for name in REQUIRED:
        path=artifact/name
        if not path.is_file() or path.stat().st_size==0:errors.append(f"核心交付缺失或为空：{name}");continue
        try:
            hashes.append(_sha256(path))
            if name==XLSX:
                if not zipfile.is_zipfile(path):raise ValueError("不是 OOXML ZIP")
                book=load_workbook(path,read_only=True,data_only=False)
                if not book.sheetnames:raise ValueError("Excel 无工作表")
                book.close()
            elif name==DOCX:
                if not zipfile.is_zipfile(path):raise ValueError("不是 OOXML ZIP")
                document=Document(path)
                if not document.paragraphs and not document.tables:raise ValueError("Word 无内容")
            else:
                text=path.read_text(encoding="utf-8")
                if "BEGIN:VCALENDAR" not in text or "END:VCALENDAR" not in text or "BEGIN:VEVENT" not in text:raise ValueError("ICS 不是可解析日历")
        except Exception as exc:errors.append(f"核心交付无法解析：{name}: {type(exc).__name__}: {exc}")
    if SOURCE_ERRORS:errors.extend(f"冻结 ICS 错误：{item}" for item in SOURCE_ERRORS)
    return not errors,errors,hashes


def _fatal(errors:list[str],hashes:list[str])->dict[str,Any]:
    results=[_criterion(key,[(False,error) for error in errors]) for key in WEIGHTS]
    results.extend([{"criterion_id":key,"score":0.0,"evidence":"未触发","method":"deterministic","raw":{}} for key in ("P001","P002")])
    return {"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":hashes or ["0"*64],"criterion_results":results,"reward":0.0,"passed":False,"hurdle_failures":["R001"],"penalties":[]}


def _recalculate(path:Path)->tuple[dict[str,list[list[Any]]]|None,str]:
    with tempfile.TemporaryDirectory(prefix="calendar-v2-recalc-") as raw:
        root=Path(raw);source=root/"source";output=root/"output";profile=root/"profile";source.mkdir();output.mkdir();copied=source/path.name;shutil.copy2(path,copied)
        env=os.environ.copy();env["HOME"]=str(root/"home");env["NO_PROXY"]="localhost,127.0.0.1,::1"
        try:completed=subprocess.run(["soffice","--headless",f"-env:UserInstallation=file://{profile}","--convert-to","xlsx","--outdir",str(output),str(copied)],stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,timeout=300,check=False,env=env)
        except Exception as exc:return None,str(exc)
        converted=output/path.name
        if completed.returncode!=0 or not converted.is_file():return None,f"退出={completed.returncode}; {completed.stdout[-400:]}"
        try:
            book=load_workbook(converted,data_only=True);values={sheet.title:[[cell.value for cell in row] for row in sheet.iter_rows()] for sheet in book.worksheets};book.close();return values,completed.stdout[-400:]
        except Exception as exc:return None,str(exc)


def _value(values:dict[str,list[list[Any]]]|None,sheet:str,row:int,column:int)->Any:
    try:return values[sheet][row-1][column-1] if values is not None else None
    except (KeyError,IndexError):return None


def _headers(sheet,count:int)->list[str]:return [V1.normalize(sheet.cell(1,column).value) for column in range(1,count+1)] if sheet else []


def _workbook(artifact:Path)->tuple[dict[str,list[tuple[bool,str]]],dict[str,Any],list[Any]]:
    findings={key:[] for key in ("R002","R003","R004","R005","R006","R007","R008","R009")};context={};danger=[]
    try:
        package=V1.inspect_package(artifact/XLSX);danger=[*package.get("macros",[]),*package.get("external",[]),*package.get("missing_targets",[])]
        book=load_workbook(artifact/XLSX,data_only=False);values,recalc_log=_recalculate(artifact/XLSX)
    except Exception as exc:
        for key in findings:findings[key].append((False,f"Excel 检查失败：{exc}"))
        return findings,context,danger
    events:list[Any]=TRUTH["events"];ids:dict[Any,str]=TRUTH["ids"];pairs:list[Any]=TRUTH["pairs"];windows:list[Any]=TRUTH["conflict_windows"];free:list[Any]=TRUTH["free"]
    for index,name in enumerate(V1.SHEETS):findings["R009"].append((index<len(book.sheetnames) and book.sheetnames[index]==name,f"工作表第 {index+1} 项错误"))

    source=book["Source_Events"] if "Source_Events" in book.sheetnames else None
    for col,target in enumerate(V1.SOURCE_HEADERS,1):findings["R002"].append((_same(source.cell(1,col).value if source else None,target),f"Source_Events 表头第 {col} 列错误"))
    source_rows=[]
    for row,event in enumerate(events,2):
        expected=(ids[event],event.owner,event.summary,event.start,event.end,event.uid,event.source_file)
        actual=(source.cell(row,1).value if source else None,source.cell(row,2).value if source else None,source.cell(row,3).value if source else None,_value(values,"Source_Events",row,4),_value(values,"Source_Events",row,5),source.cell(row,7).value if source else None,source.cell(row,8).value if source else None)
        for pos,(got,target) in enumerate(zip(actual,expected,strict=True)):
            ok=_time(got,target) if pos in (3,4) else _same(got,target);findings["R002"].append((ok,f"Source_Events 第 {row} 行字段 {pos+1} 错误"))
        internal=sum(1 for other in events if other.owner==event.owner and other!=event and V1.overlap(event,other));cross=sum(1 for other in events if other.owner!=event.owner and V1.overlap(event,other));duration=(event.end-event.start).total_seconds()/60
        formula_specs=((6,(f"E{row}-D{row}","1440"),duration),(9,("SUMPRODUCT",f"B2:B11=B{row}",f"A2:A11<>A{row}",f"D2:D11<E{row}",f"E2:E11>D{row}"),internal),(10,("SUMPRODUCT",f"B2:B11<>B{row}",f"D2:D11<E{row}",f"E2:E11>D{row}"),cross))
        for col,tokens,target in formula_specs:findings["R003"].extend([(_formula(source.cell(row,col).value if source else None,*tokens),f"Source_Events!{chr(64+col)}{row} 公式错误"),(_same(_value(values,"Source_Events",row,col),target),f"Source_Events!{chr(64+col)}{row} 重算错误：{recalc_log}")])
        findings["R004"].append((_same(_value(values,"Source_Events",row,9),0),f"{ids[event]} 被误算内部冲突"));source_rows.append(actual)
    findings["R004"].append((len(TRUTH["internal"])==0,"来源独立重算的内部冲突不为 0"));context["source_rows"]=source_rows

    pair_sheet=book["Pairwise_Conflicts"] if "Pairwise_Conflicts" in book.sheetnames else None
    for col,target in enumerate(V1.PAIR_HEADERS,1):findings["R005"].append((_same(pair_sheet.cell(1,col).value if pair_sheet else None,target),f"Pairwise_Conflicts 表头第 {col} 列错误"))
    pair_rows=[]
    for row,(bob,tom,start,end) in enumerate(pairs,2):
        expected=(f"P{row-1:02d}",ids[bob],bob.summary,ids[tom],tom.summary,start,end,(end-start).total_seconds()/60,"Yes" if bob.summary==tom.summary else "No")
        actual=tuple(_value(values,"Pairwise_Conflicts",row,col) if col in (6,7,8,9) else (pair_sheet.cell(row,col).value if pair_sheet else None) for col in range(1,10))
        for pos,(got,target) in enumerate(zip(actual,expected,strict=True)):
            ok=_time(got,target) if pos in (5,6) else _same(got,target);findings["R005"].append((ok,f"Pairwise_Conflicts 第 {row} 行第 {pos+1} 列错误"))
        findings["R005"].extend([(_formula(pair_sheet.cell(row,8).value if pair_sheet else None,f"G{row}-F{row}","1440"),f"Pairwise_Conflicts!H{row} 公式错误"),(_formula(pair_sheet.cell(row,9).value if pair_sheet else None,"IF",f"C{row}=E{row}"),f"Pairwise_Conflicts!I{row} 公式错误")]);pair_rows.append(actual)
    context["pair_rows"]=pair_rows

    window_sheet=book["Conflict_Windows"] if "Conflict_Windows" in book.sheetnames else None
    for col,target in enumerate(V1.WINDOW_HEADERS,1):findings["R006"].append((_same(window_sheet.cell(1,col).value if window_sheet else None,target),f"Conflict_Windows 表头第 {col} 列错误"))
    window_rows=[]
    for row,(start,end) in enumerate(windows,2):
        bob_names="; ".join(sorted({event.summary for event in events if event.owner=="Bob" and event.start<end and event.end>start}));tom_names="; ".join(sorted({event.summary for event in events if event.owner=="Tom" and event.start<end and event.end>start}));count=sum(1 for _,_,pstart,pend in pairs if pstart<end and pend>start)
        expected=(f"W{row-1:02d}",start,end,(end-start).total_seconds()/60,bob_names,tom_names,count);actual=tuple(_value(values,"Conflict_Windows",row,col) if col in (2,3,4,7) else (window_sheet.cell(row,col).value if window_sheet else None) for col in range(1,8))
        for pos,(got,target) in enumerate(zip(actual,expected,strict=True)):
            ok=_time(got,target) if pos in (1,2) else _same(got,target);findings["R006"].append((ok,f"Conflict_Windows 第 {row} 行第 {pos+1} 列错误"))
        findings["R006"].extend([(_formula(window_sheet.cell(row,4).value if window_sheet else None,f"C{row}-B{row}","1440"),f"Conflict_Windows!D{row} 公式错误"),(_formula(window_sheet.cell(row,7).value if window_sheet else None,"COUNTIFS","PAIRWISE_CONFLICTS",f"C{row}",f"B{row}"),f"Conflict_Windows!G{row} 公式错误")]);window_rows.append(actual)
    context["window_rows"]=window_rows

    free_sheet=book["Joint_Free_Windows"] if "Joint_Free_Windows" in book.sheetnames else None
    for col,target in enumerate(V1.FREE_HEADERS,1):findings["R007"].append((_same(free_sheet.cell(1,col).value if free_sheet else None,target),f"Joint_Free_Windows 表头第 {col} 列错误"))
    free_rows=[]
    for row,(start,end) in enumerate(free,2):
        minutes=(end-start).total_seconds()/60;expected=(f"F{row-1:02d}","2024-05-01",start,end,minutes,"Yes" if minutes==90 else "No","Recommended longest joint window" if minutes==90 else "Available alternative");actual=tuple(_value(values,"Joint_Free_Windows",row,col) if col in (3,4,5,6,7) else (free_sheet.cell(row,col).value if free_sheet else None) for col in range(1,8))
        for pos,(got,target) in enumerate(zip(actual,expected,strict=True)):
            ok=_time(got,target) if pos in (2,3) else _same(got,target);findings["R007"].append((ok,f"Joint_Free_Windows 第 {row} 行第 {pos+1} 列错误"))
        findings["R007"].extend([(_formula(free_sheet.cell(row,5).value if free_sheet else None,f"D{row}-C{row}","1440"),f"Joint_Free_Windows!E{row} 公式错误"),(_formula(free_sheet.cell(row,6).value if free_sheet else None,"IF",f"E{row}=MAX(E2:E5)"),f"Joint_Free_Windows!F{row} 公式错误"),(_formula(free_sheet.cell(row,7).value if free_sheet else None,"IF",f"F{row}=\"YES\""),f"Joint_Free_Windows!G{row} 公式错误")]);free_rows.append(actual)
    context["free_rows"]=free_rows

    summary=book["Summary"] if "Summary" in book.sheetnames else None;expected_values=[10,5,5,0,5,4,420,4,160,90,2,2,10]
    formula_tokens=[("COUNTA","SOURCE_EVENTS!A2:A11"),("COUNTIF","SOURCE_EVENTS!B2:B11","BOB"),("COUNTIF","SOURCE_EVENTS!B2:B11","TOM"),("SUM","SOURCE_EVENTS!I2:I11","/2"),("COUNTA","PAIRWISE_CONFLICTS!A2:A6"),("COUNTA","CONFLICT_WINDOWS!A2:A5"),("SUM","CONFLICT_WINDOWS!D2:D5"),("COUNTA","JOINT_FREE_WINDOWS!A2:A5"),("SUM","JOINT_FREE_WINDOWS!E2:E5"),("MAX","JOINT_FREE_WINDOWS!E2:E5"),("COUNTIF","PAIRWISE_CONFLICTS!I2:I6","YES"),("SUMPRODUCT","COUNTIF","SOURCE_EVENTS!H2:H11"),("COUNTA","SOURCE_EVENTS!A2:A11")]
    summary_values={}
    for row,(label,target,tokens) in enumerate(zip(V1.SUMMARY_LABELS,expected_values,formula_tokens,strict=True),2):
        findings["R008"].extend([(_same(summary.cell(row,1).value if summary else None,label),f"Summary!A{row} 标签错误"),(_formula(summary.cell(row,2).value if summary else None,*tokens),f"Summary!B{row} 公式错误"),(_same(_value(values,"Summary",row,2),target),f"Summary!B{row} 重算错误")]);summary_values[label]=_value(values,"Summary",row,2)
    context["summary"]=summary_values

    tables={"Source_Events":"CalendarSourceEventsTable","Pairwise_Conflicts":"CalendarPairwiseConflictsTable","Conflict_Windows":"CalendarConflictWindowsTable","Joint_Free_Windows":"CalendarJointFreeWindowsTable","Summary":"CalendarAuditSummaryTable"}
    for sheet_name,table_name in tables.items():
        sheet=book[sheet_name] if sheet_name in book.sheetnames else None;findings["R009"].extend([(sheet is not None and set(sheet.tables)=={table_name},f"{sheet_name} 原生表错误"),(sheet is not None and sheet.freeze_panes=="A2",f"{sheet_name} 未冻结首行")])
    findings["R009"].extend([(summary is not None and len(summary._charts)==1,"Summary 缺可编辑时长图"),(sum(len(sheet.conditional_formatting) for sheet in book.worksheets)>=6,"条件格式少于 6 处"),(not any(len(sheet._images) for sheet in book.worksheets),"Excel 用图片代替数据"),(not danger,f"OOXML 含危险/断裂对象：{danger}")])
    book.close();return findings,context,danger


def _calendar(artifact:Path)->tuple[list[tuple[bool,str]],dict[str,Any],list[str]]:
    findings=[];events,meta=V1.parse_ics_events(artifact/ICS,output=True);source_events:list[Any]=TRUTH["events"]
    for key,target in (("VERSION","2.0"),("CALSCALE","GREGORIAN"),("METHOD","PUBLISH"),("X-WR-CALNAME","Bob and Tom Merged Schedule")):findings.append((_same(meta["calendar"].get(key),target),f"VCALENDAR {key} 错误"))
    findings.append((not meta["errors"],f"ICS 格式/属性错误：{meta['errors']}"));findings.append((len(events)==10,f"ICS 事件数错误：{len(events)}"));findings.append((len({event.uid for event in events})==10,"ICS UID 不唯一"))
    for index,target in enumerate(source_events):
        got=events[index] if index<len(events) else None
        for attr in ("uid","owner","source_file","summary"):
            findings.append((got is not None and _same(getattr(got,attr),getattr(target,attr)),f"ICS 第 {index+1} 事件 {attr} 错误"))
        findings.extend([(got is not None and _time(got.start,target.start),f"ICS 第 {index+1} 事件开始时间错误"),(got is not None and _time(got.end,target.end),f"ICS 第 {index+1} 事件结束时间错误")])
    unsupported=[str(item) for item in meta["errors"] if "unsupported_property" in str(item)]
    return findings,{"events":events,"count":len(events),"meta":meta},unsupported


def _rows(table)->list[list[str]]:return [[cell.text for cell in row.cells] for row in table.rows]


def _word(artifact:Path)->tuple[dict[str,list[tuple[bool,str]]],dict[str,Any],list[str]]:
    findings={"R011":[],"R012":[]};document=Document(artifact/DOCX);text=V1.word_text(document);lower=_norm(text);headings=[V1.normalize(p.text) for p in document.paragraphs if str(p.style.name).startswith("Heading")];events:list[Any]=TRUTH["events"];ids:dict[Any,str]=TRUTH["ids"];pairs:list[Any]=TRUTH["pairs"];windows:list[Any]=TRUTH["conflict_windows"];free:list[Any]=TRUTH["free"]
    findings["R011"].append(("Calendar Conflict Audit Report" in text,"Word 标题错误"))
    for heading in ("Scope and Time Semantics","Source Event Reconciliation","Conflict Findings","Joint Availability","Merged Calendar Control","Audit Conclusion"):findings["R011"].append((heading in headings,f"Word 缺章节 {heading}"))
    for token in ("half-open intervals","touching endpoints are not conflicts","without local-time conversion","X-EVOCOWORK-OWNER","X-EVOCOWORK-SOURCE"):findings["R011"].append((token.casefold() in lower,f"Word 边界说明缺 {token}"))
    findings["R011"].extend([(len(document.tables)==5,f"Word 原生表数错误：{len(document.tables)}"),(not document.inline_shapes,"Word 用图片代替表格")])
    expected=[]
    expected.append([["ID","Owner","Summary","Start","End","Min","Source"]]+[[ids[event],event.owner,event.summary,event.start.strftime("%H:%M"),event.end.strftime("%H:%M"),str(int((event.end-event.start).total_seconds()/60)),event.source_file] for event in events])
    expected.append([["Pair","Bob event","Tom event","Start","End","Min","Same title"]]+[[f"P{i:02d}",f"{ids[bob]} {bob.summary}",f"{ids[tom]} {tom.summary}",start.strftime("%H:%M"),end.strftime("%H:%M"),str(int((end-start).total_seconds()/60)),"Yes" if bob.summary==tom.summary else "No"] for i,(bob,tom,start,end) in enumerate(pairs,1)])
    window_rows=[]
    for i,(start,end) in enumerate(windows,1):
        bob="; ".join(sorted({event.summary for event in events if event.owner=="Bob" and event.start<end and event.end>start}));tom="; ".join(sorted({event.summary for event in events if event.owner=="Tom" and event.start<end and event.end>start}));count=sum(1 for _,_,pstart,pend in pairs if pstart<end and pend>start);window_rows.append([f"W{i:02d}",start.strftime("%H:%M"),end.strftime("%H:%M"),str(int((end-start).total_seconds()/60)),bob,tom,str(count)])
    expected.append([["Window","Start","End","Min","Bob active","Tom active","Pairs"]]+window_rows)
    expected.append([["Window","Date","Start UTC","End UTC","Minutes","Longest"]]+[[f"F{i:02d}","2024-05-01",start.strftime("%H:%M"),end.strftime("%H:%M"),str(int((end-start).total_seconds()/60)),"Yes" if int((end-start).total_seconds()/60)==90 else "No"] for i,(start,end) in enumerate(free,1)])
    expected.append([["Control","Expected","Result"],["Source files","Bob.ics + Tom.ics","2 reconciled"],["VEVENT count","5 + 5","10 retained"],["UID/time/title changes","None","None"],["Owner/source extensions","Present on every event","10/10"]])
    actual_tables=[_rows(table) for table in document.tables]
    for ti,table_rows in enumerate(expected):
        for ri,row in enumerate(table_rows):
            for ci,target in enumerate(row):
                actual=actual_tables[ti][ri][ci] if ti<len(actual_tables) and ri<len(actual_tables[ti]) and ci<len(actual_tables[ti][ri]) else None;findings["R012"].append((_same(actual,target),f"Word 表 {ti+1} 第 {ri+1} 行第 {ci+1} 列错误"))
    for token in ("10 VEVENT","five for Bob","five for Tom","0 internal conflict pairs","five event pairs","four merged conflict windows","420 minutes","four windows totaling 160 minutes","14:30-16:00 UTC","90 minutes","class at 16:00-18:00 UTC","dinner at 18:30-19:30 UTC","No timezone, attendee, location, priority or additional meeting is inferred"):findings["R012"].append((token.casefold() in lower,f"Word 缺核心事实 {token}"))
    unsupported=V1.unsupported_claims(text);return findings,{"text":text,"lower":lower,"tables":actual_tables},unsupported


def evaluate(artifact:Path,scoring_version:str|None=None)->dict[str,Any]:
    if (scoring_version or os.environ.get("SCORING_VERSION","v2")).casefold()=="v1":return V1.evaluate(artifact)
    core_ok,errors,hashes=_core(artifact)
    if not core_ok:return _fatal(errors,hashes)
    criteria={key:[] for key in WEIGHTS};criteria["R001"].append((True,"三份核心交付和两份冻结 ICS 可解析"));workbook,wctx,danger=_workbook(artifact)
    for key,items in workbook.items():criteria[key].extend(items)
    calendar,cctx,ics_unsupported=_calendar(artifact);criteria["R010"].extend(calendar);word,wordctx,word_unsupported=_word(artifact)
    for key,items in word.items():criteria[key].extend(items)
    summary=wctx.get("summary",{});source_rows=wctx.get("source_rows",[]);pairs=wctx.get("pair_rows",[]);windows=wctx.get("window_rows",[]);free=wctx.get("free_rows",[]);lower=wordctx.get("lower","")
    consistency=[
        (_same(summary.get("Total_Events"),len(source_rows)) and _same(summary.get("Total_Events"),cctx.get("count")) and f"{int(summary.get('Total_Events') or 0)} vevent" in lower,"三产物事件总数不一致"),
        (_same(summary.get("Bob_Events"),sum(_same(row[1],"Bob") for row in source_rows)) and "five for bob" in lower,"Bob 事件数不一致"),
        (_same(summary.get("Tom_Events"),sum(_same(row[1],"Tom") for row in source_rows)) and "five for tom" in lower,"Tom 事件数不一致"),
        (_same(summary.get("Cross_Owner_Pairwise_Conflicts"),len(pairs)) and "five event pairs" in lower,"pairwise 冲突数不一致"),
        (_same(summary.get("Merged_Conflict_Windows"),len(windows)) and "four merged conflict windows" in lower,"合并冲突窗口数不一致"),
        (_same(summary.get("Total_Conflict_Window_Minutes"),sum(float(row[3] or 0) for row in windows)) and "420 minutes" in lower,"冲突窗口分钟数不一致"),
        (_same(summary.get("Joint_Free_Windows"),len(free)) and "four windows" in lower,"空闲窗口数不一致"),
        (_same(summary.get("Total_Joint_Free_Minutes"),sum(float(row[4] or 0) for row in free)) and "160 minutes" in lower,"空闲分钟数不一致"),
        (_same(summary.get("Merged_ICS_Events"),cctx.get("count")) and "10 retained" in lower,"合并 ICS 控制数不一致"),
    ];criteria["R013"].extend(consistency)
    render=V1.render_office(artifact);criteria["R014"].extend([(not render.get("errors"),f"LibreOffice 渲染错误：{render.get('errors')}"),(render.get("xlsx_pages")==5,f"Excel 页数错误：{render.get('xlsx_pages')}"),(render.get("docx_pages")==3,f"Word 页数错误：{render.get('docx_pages')}")])
    extra=sorted(path.name for path in artifact.iterdir() if path.name not in REQUIRED);suspicious=V1.suspicious_content(artifact);unsupported=word_unsupported+ics_unsupported;p1_score=min(1.0,len(unsupported)/2.0);p2_bad=bool(extra or danger or suspicious)
    positive=[_criterion(key,criteria[key]) for key in WEIGHTS];base=sum(WEIGHTS[item["criterion_id"]]*float(item["score"]) for item in positive)/100.0;applied=[]
    if p1_score:applied.append({"id":"P001","amount":round(.12*p1_score,6),"reason":"编造时区/参会人/地点/优先级/新会议"})
    if p2_bad:applied.append({"id":"P002","amount":.04,"reason":f"额外交付或危险对象：extra={extra}"})
    reward=round(max(0.0,base-sum(item["amount"] for item in applied)),6);positive.extend([{"criterion_id":"P001","score":round(p1_score,6),"evidence":"未触发" if not p1_score else "检测到无来源断言/属性","method":"deterministic","raw":{"unsupported":unsupported}},{"criterion_id":"P002","score":float(p2_bad),"evidence":"未触发" if not p2_bad else "检测到额外交付/危险对象","method":"deterministic","raw":{"extra":extra,"danger":danger,"suspicious":suspicious}}])
    return {"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":hashes,"criterion_results":positive,"base_score":round(base,6),"reward":reward,"passed":reward>=.8,"hurdle_failures":[],"penalties":applied}


def main()->None:
    parser=argparse.ArgumentParser();parser.add_argument("--artifact-dir",type=Path,required=True);parser.add_argument("--result",type=Path,required=True);parser.add_argument("--reward-file",type=Path,required=True);parser.add_argument("--scoring-version",choices=("v1","v2"),default=os.environ.get("SCORING_VERSION","v2"));args=parser.parse_args()
    try:result=evaluate(args.artifact_dir,args.scoring_version)
    except Exception as exc:result={"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"reward":0.0,"passed":False,"fatal_error":f"{type(exc).__name__}: {exc}","hurdle_failures":["R001"]}
    args.result.parent.mkdir(parents=True,exist_ok=True);args.reward_file.parent.mkdir(parents=True,exist_ok=True);args.result.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n",encoding="utf-8");args.reward_file.write_text(f"{float(result['reward']):.6f}\n",encoding="utf-8");print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=="__main__":main()
