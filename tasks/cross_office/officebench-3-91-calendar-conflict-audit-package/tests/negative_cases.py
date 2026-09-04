#!/usr/bin/env python3
"""日历冲突审计包的 v2 任务专属分层负例。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import re
import shutil
import sys
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Callable

from docx import Document
from openpyxl import Workbook, load_workbook


ROOT=Path(__file__).resolve().parents[1];GOLDEN=ROOT/"solution"/"Golden";INPUTS=ROOT/"environment"/"input_files"
REQUIRED=("Calendar_Conflict_Audit.xlsx","Merged_Schedule.ics","Calendar_Conflict_Report.docx");Mutator=Callable[[Path],None]


def _load_judge():
    path=Path(__file__).with_name("judge.py");spec=importlib.util.spec_from_file_location("calendar_v2_judge_for_negatives",path)
    if spec is None or spec.loader is None:raise RuntimeError("无法载入 v2 Judge")
    module=importlib.util.module_from_spec(spec);sys.modules[spec.name]=module;spec.loader.exec_module(module);return module


JUDGE=_load_judge()


def _copy(target:Path)->None:
    target.mkdir(parents=True,exist_ok=True)
    for name in REQUIRED:shutil.copy2(GOLDEN/name,target/name)


def _excel(sheet:str,cell:str,value:object)->Mutator:
    def apply(root:Path)->None:
        path=root/REQUIRED[0];book=load_workbook(path);book[sheet][cell]=value;book.save(path);book.close()
    return apply


def _remove_chart(root:Path)->None:
    path=root/REQUIRED[0];book=load_workbook(path);book["Summary"]._charts=[];book.save(path);book.close()


def _remove_freeze(root:Path)->None:
    path=root/REQUIRED[0];book=load_workbook(path);book["Joint_Free_Windows"].freeze_panes=None;book.save(path);book.close()


def _remove_table(root:Path)->None:
    path=root/REQUIRED[0];book=load_workbook(path);del book["Pairwise_Conflicts"].tables["CalendarPairwiseConflictsTable"];book.save(path);book.close()


def _remove_conditional(root:Path)->None:
    path=root/REQUIRED[0];book=load_workbook(path);book["Source_Events"].conditional_formatting._cf_rules.clear();book.save(path);book.close()


def _ics_block(uid:str,old:str,new:str)->Mutator:
    def apply(root:Path)->None:
        path=root/REQUIRED[1];text=path.read_bytes().decode("utf-8");blocks=re.findall(r"BEGIN:VEVENT\r\n(?:(?!END:VEVENT).)*?END:VEVENT\r\n",text,flags=re.DOTALL);target=next((block for block in blocks if f"UID:{uid}" in block),None)
        if target is None or old not in target:raise AssertionError(f"ICS 未找到 {uid} / {old}")
        path.write_bytes(text.replace(target,target.replace(old,new,1),1).encode("utf-8"))
    return apply


def _ics_drop(root:Path)->None:
    path=root/REQUIRED[1];text=path.read_bytes().decode("utf-8");updated,count=re.subn(r"BEGIN:VEVENT\r\n(?:(?!END:VEVENT).)*?UID:13e17887-2bbc-42a2-a881-5aabcdfda201@13e1\.org(?:(?!END:VEVENT).)*?END:VEVENT\r\n","",text,count=1,flags=re.DOTALL);assert count==1;path.write_bytes(updated.encode("utf-8"))


def _ics_add(root:Path)->None:
    path=root/REQUIRED[1];text=path.read_bytes().decode("utf-8");extra="BEGIN:VEVENT\r\nUID:invented@evocowork\r\nDTSTAMP:20260901T063000Z\r\nDTSTART:20240501T150000Z\r\nDTEND:20240501T153000Z\r\nSUMMARY:invented meeting\r\nX-EVOCOWORK-OWNER:Bob\r\nX-EVOCOWORK-SOURCE:Bob.ics\r\nEND:VEVENT\r\n";path.write_bytes(text.replace("END:VCALENDAR\r\n",extra+"END:VCALENDAR\r\n").encode("utf-8"))


def _ics_lf(root:Path)->None:
    path=root/REQUIRED[1];path.write_bytes(path.read_bytes().replace(b"\r\n",b"\n"))


def _word_replace(old:str,new:str)->Mutator:
    def apply(root:Path)->None:
        path=root/REQUIRED[2];document=Document(path);locations=list(document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:locations.extend(cell.paragraphs)
        for paragraph in locations:
            if old in paragraph.text:paragraph.text=paragraph.text.replace(old,new,1);document.save(path);return
        raise AssertionError(f"Word 未找到 {old}")
    return apply


def _word_table(t:int,r:int,c:int,value:str)->Mutator:
    def apply(root:Path)->None:
        path=root/REQUIRED[2];document=Document(path);document.tables[t].rows[r].cells[c].text=value;document.save(path)
    return apply


def _word_remove_table(root:Path)->None:
    path=root/REQUIRED[2];document=Document(path);table=document.tables[1];table._element.getparent().remove(table._element);document.save(path)


def _word_add(text:str)->Mutator:
    def apply(root:Path)->None:
        path=root/REQUIRED[2];document=Document(path);document.add_paragraph(text);document.save(path)
    return apply


def _moderate_analysis_conflict(root:Path)->None:
    path=root/REQUIRED[0];book=load_workbook(path)
    for sheet_name,start,end in (("Pairwise_Conflicts",2,6),("Conflict_Windows",2,5),("Joint_Free_Windows",2,5)):
        sheet=book[sheet_name]
        for row in range(start,end+1):
            for col in range(1,sheet.max_column+1):sheet.cell(row,col).value=0
    summary=book["Summary"]
    for row in range(5,15):summary.cell(row,2).value=0
    book.save(path);book.close()
    doc=root/REQUIRED[2];document=Document(doc)
    for index in (1,2,3):
        for row in document.tables[index].rows[1:]:
            for cell in row.cells:cell.text="0"
    document.paragraphs[7].text="Across owners, no event pairs overlap and there are no conflict windows."
    document.paragraphs[10].text="No joint availability was calculated."
    document.save(doc)


def _severe_valid_incomplete(root:Path)->None:
    book=Workbook();book.remove(book.active)
    headers={"Source_Events":JUDGE.V1.SOURCE_HEADERS,"Pairwise_Conflicts":JUDGE.V1.PAIR_HEADERS,"Conflict_Windows":JUDGE.V1.WINDOW_HEADERS,"Joint_Free_Windows":JUDGE.V1.FREE_HEADERS,"Summary":["Metric","Value","Evidence"]}
    for name,row in headers.items():sheet=book.create_sheet(name);sheet.append(row)
    book.save(root/REQUIRED[0]);book.close()
    event=JUDGE.TRUTH["events"][0];ics=("BEGIN:VCALENDAR\r\nVERSION:2.0\r\nBEGIN:VEVENT\r\n"+f"UID:{event.uid}\r\nDTSTART:{JUDGE.V1.format_utc(event.start)}\r\nDTEND:{JUDGE.V1.format_utc(event.end)}\r\nSUMMARY:{event.summary}\r\nX-EVOCOWORK-OWNER:{event.owner}\r\nX-EVOCOWORK-SOURCE:{event.source_file}\r\nEND:VEVENT\r\nEND:VCALENDAR\r\n");(root/REQUIRED[1]).write_bytes(ics.encode("utf-8"))
    document=Document();document.add_paragraph("Calendar Conflict Audit Report");document.add_heading("Scope and Time Semantics",level=1);document.add_paragraph("Audit incomplete.");document.save(root/REQUIRED[2])


def _input_only(root:Path)->None:
    for path in list(root.iterdir()):
        if path.is_file():path.unlink()
    for source in INPUTS.iterdir():
        if source.is_file():shutil.copy2(source,root/source.name)


PARTIALS:dict[str,tuple[Mutator,bool]]={
    "source_event_id_wrong":(_excel("Source_Events","A2","E99"),False),"source_owner_wrong":(_excel("Source_Events","B2","Alice"),False),"source_summary_wrong":(_excel("Source_Events","C2","sleep"),False),"source_start_wrong":(_excel("Source_Events","D2",datetime(2024,5,1,8,30)),False),"source_end_wrong":(_excel("Source_Events","E2",datetime(2024,5,1,10,30)),False),"source_uid_wrong":(_excel("Source_Events","G2","wrong@uid"),False),"source_file_wrong":(_excel("Source_Events","H2","Tom.ics"),False),
    "duration_flattened":(_excel("Source_Events","F2",180),False),"duration_wrong_formula":(_excel("Source_Events","F2","=(E2-D2)*60"),False),"internal_conflict_flattened":(_excel("Source_Events","I2",0),False),"internal_conflict_fabricated":(_excel("Source_Events","I2","=1"),False),"cross_conflict_flattened":(_excel("Source_Events","J2",1),False),"cross_conflict_wrong_formula":(_excel("Source_Events","J2","=0"),False),
    "pair_id_wrong":(_excel("Pairwise_Conflicts","A2","P99"),False),"pair_bob_id_wrong":(_excel("Pairwise_Conflicts","B2","E03"),False),"pair_tom_summary_wrong":(_excel("Pairwise_Conflicts","E2","meeting"),False),"pair_start_wrong":(_excel("Pairwise_Conflicts","F2",datetime(2024,5,1,8,30)),False),"pair_end_wrong":(_excel("Pairwise_Conflicts","G2",datetime(2024,5,1,10,30)),False),"pair_overlap_flattened":(_excel("Pairwise_Conflicts","H2",180),False),"pair_overlap_wrong_formula":(_excel("Pairwise_Conflicts","H2","=(G2-F2)*60"),False),"pair_same_title_flattened":(_excel("Pairwise_Conflicts","I5","Yes"),False),"pair_same_title_wrong_formula":(_excel("Pairwise_Conflicts","I5","=IF(C5=C5,\"Yes\",\"No\")"),False),
    "window_id_wrong":(_excel("Conflict_Windows","A3","W09"),False),"window_end_wrong":(_excel("Conflict_Windows","C3",datetime(2024,5,1,13,0)),False),"window_duration_flattened":(_excel("Conflict_Windows","D3",60),False),"window_active_bob_wrong":(_excel("Conflict_Windows","E3","lunch"),False),"window_pair_count_wrong_formula":(_excel("Conflict_Windows","G3","=1"),False),
    "free_id_wrong":(_excel("Joint_Free_Windows","A3","F09"),False),"free_start_wrong":(_excel("Joint_Free_Windows","C3",datetime(2024,5,1,15,0)),False),"free_duration_flattened":(_excel("Joint_Free_Windows","E3",90),False),"free_longest_wrong_formula":(_excel("Joint_Free_Windows","F3","=\"No\""),False),"free_recommendation_flattened":(_excel("Joint_Free_Windows","G3","Recommended longest joint window"),False),
    "summary_label_wrong":(_excel("Summary","A2","Event_Total"),False),"summary_event_count_flattened":(_excel("Summary","B2",10),False),"summary_internal_wrong_formula":(_excel("Summary","B5","=1"),False),"summary_conflict_minutes_wrong":(_excel("Summary","B8","=400"),False),"summary_free_minutes_flattened":(_excel("Summary","B10",160),False),"summary_ics_count_wrong":(_excel("Summary","B14","=9"),False),
    "native_chart_removed":(_remove_chart,False),"native_freeze_removed":(_remove_freeze,False),"native_table_removed":(_remove_table,False),"native_conditional_format_removed":(_remove_conditional,False),
    "ics_event_dropped":(_ics_drop,False),"ics_time_changed":(_ics_block("2fc598fd-361b-46c0-8b12-426c4d7939af@2fc5.org","DTSTART:20240501T160000Z","DTSTART:20240501T163000Z"),False),"ics_summary_changed":(_ics_block("f825c769-a892-426d-8aff-31e2801edfe2@f825.org","SUMMARY:dinner","SUMMARY:supper"),False),"ics_owner_changed":(_ics_block("7cf33664-81fa-44da-860a-6076c22bbf70@7cf3.org","X-EVOCOWORK-OWNER:Bob","X-EVOCOWORK-OWNER:Tom"),False),"ics_source_changed":(_ics_block("7cf33664-81fa-44da-860a-6076c22bbf70@7cf3.org","X-EVOCOWORK-SOURCE:Bob.ics","X-EVOCOWORK-SOURCE:Tom.ics"),False),"ics_event_added":(_ics_add,False),"ics_lf_not_crlf":(_ics_lf,False),
    "word_heading_wrong":(_word_replace("Joint Availability","Availability Review"),False),"word_event_wrong":(_word_table(0,1,2,"sleep"),False),"word_pair_minutes_wrong":(_word_table(1,1,5,"170"),False),"word_window_active_wrong":(_word_table(2,2,4,"lunch"),False),"word_free_longest_wrong":(_word_table(3,2,5,"No"),False),"word_ics_control_wrong":(_word_table(4,2,2,"9 retained"),False),"word_conflict_total_wrong":(_word_replace("420 minutes","400 minutes"),False),"word_timezone_boundary_missing":(_word_replace("without local-time conversion","after conversion"),False),"word_pair_table_missing":(_word_remove_table,False),"word_invented_timezone_meeting":(_word_add("The schedule was converted to EST and a new meeting was scheduled."),False),
    "unexpected_debug_file":(lambda root:(root/"debug.txt").write_text("temporary\n",encoding="utf-8"),False),"moderate_analysis_domain_conflict":(_moderate_analysis_conflict,True),"severe_valid_but_incomplete_package":(_severe_valid_incomplete,True),
}


def main()->None:
    parser=argparse.ArgumentParser();parser.add_argument("--report",type=Path);args=parser.parse_args();golden_recalc,_=JUDGE._recalculate(GOLDEN/REQUIRED[0]);golden_render=JUDGE.V1.render_office(GOLDEN);original_recalc,original_render=JUDGE._recalculate,JUDGE.V1.render_office
    assert JUDGE.evaluate(GOLDEN)["reward"]==1.0
    def fast_recalc(path:Path):
        book=load_workbook(path,data_only=False);values={}
        for sheet in book.worksheets:
            rows=[]
            for row_index,row in enumerate(sheet.iter_rows(),1):
                current=[]
                for column_index,cell in enumerate(row,1):
                    value=cell.value
                    if isinstance(value,str) and value.startswith("="):value=JUDGE._value(golden_recalc,sheet.title,row_index,column_index)
                    current.append(value)
                rows.append(current)
            values[sheet.title]=rows
        book.close();return values,"任务专属 Golden 公式缓存值"
    def evaluate_case(case:Path,real:bool)->float:
        if real:return float(JUDGE.evaluate(case)["reward"])
        JUDGE._recalculate=fast_recalc;JUDGE.V1.render_office=lambda _artifact:golden_render
        try:return float(JUDGE.evaluate(case)["reward"])
        finally:JUDGE._recalculate,JUDGE.V1.render_office=original_recalc,original_render
    records=[]
    with tempfile.TemporaryDirectory(prefix="calendar-v2-negatives-") as raw:
        base=Path(raw)
        for name in REQUIRED:
            for mode in ("missing","empty","corrupt"):
                case=base/f"fatal-{mode}-{name}";_copy(case);target=case/name
                if mode=="missing":target.unlink()
                elif mode=="empty":target.write_bytes(b"")
                else:target.write_bytes(b"not-a-valid-task-artifact")
                reward=evaluate_case(case,False);assert reward==0,(mode,name,reward);records.append({"case":f"{mode}:{name}","layer":"fatal_zero","reward":reward})
        empty=base/"fatal-empty";empty.mkdir();reward=evaluate_case(empty,False);assert reward==0;records.append({"case":"empty-output","layer":"fatal_zero","reward":reward})
        input_only=base/"fatal-input-only";_copy(input_only);_input_only(input_only);reward=evaluate_case(input_only,False);assert reward==0;records.append({"case":"unmodified-input-only","layer":"fatal_zero","reward":reward})
        for name,(mutator,real) in PARTIALS.items():
            case=base/f"partial-{name}";_copy(case);mutator(case);reward=evaluate_case(case,real);assert 0<reward<1,(name,reward);records.append({"case":name,"layer":"partial","reward":reward})
    partial=[float(record["reward"]) for record in records if record["layer"]=="partial"];by_name={str(record["case"]):float(record["reward"]) for record in records};assert by_name["moderate_analysis_domain_conflict"]<.8;assert 0<by_name["severe_valid_but_incomplete_package"]<by_name["moderate_analysis_domain_conflict"]
    report={"total":len(records),"fatal_zero":sum(record["layer"]=="fatal_zero" for record in records),"partial_continuous":len(partial),"partial_reward_range":[min(partial),max(partial)],"moderate_analysis_domain_conflict":by_name["moderate_analysis_domain_conflict"],"severe_valid_but_incomplete":by_name["severe_valid_but_incomplete_package"],"cases":records}
    if args.report:args.report.parent.mkdir(parents=True,exist_ok=True);args.report.write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=="__main__":main()
