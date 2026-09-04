#!/usr/bin/env python3
"""OfficeVal 059 的 v2 连续评分入口；原严格口径保存在 judge_v1.py。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import os
import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path, PurePosixPath
from typing import Any, Callable

from lxml import etree
from pptx import Presentation


TASK_ID = "cowork-evo/officeval-059-ppt-link-in-word-application"
RUBRIC_VERSION = "omegause-officeval-059-adaptation-v0.2"
JUDGE_VERSION = "officeval-059-criterion-partial-v2.0"
DOC_NAME = "青梧谷灵MR共创计划_项目申报书_已链接PPT.docx"
SOURCE_DOC_NAME = "青梧谷灵MR共创计划_项目申报书.docx"
PPT_NAME = "青梧谷灵MR共创计划.pptx"
REQUIRED = (DOC_NAME, PPT_NAME)
SOURCE_HASHES = {
    SOURCE_DOC_NAME: "57424ad5fd58bda00674e69fc1643d22df125aaade34fbc6a4cced2e92ebc785",
    PPT_NAME: "6c0bffcfa7b1cd6d0e4d5f776218aa82880d310f7b55171d3f0cc04029770fa5",
}
POSITIVE_WEIGHTS = {
    "R001": 6,
    "R002": 14,
    "R003": 15,
    "R004": 12,
    "R005": 9,
    "R006": 14,
    "R007": 12,
    "R008": 10,
    "R009": 8,
}
PENALTY_WEIGHTS = {"P001": -12, "P002": -4}
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"w": W_NS, "r": R_NS, "rel": REL_NS}


def _load_v1():
    path = Path(__file__).with_name("judge_v1.py")
    spec = importlib.util.spec_from_file_location("officeval_059_judge_v1", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v1 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


V1 = _load_v1()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _zip_payloads(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as archive:
        bad = archive.testzip()
        if bad:
            raise ValueError(f"OOXML 部件损坏：{bad}")
        return {name: archive.read(name) for name in archive.namelist()}


def _parse_docx(path: Path) -> tuple[dict[str, bytes], etree._Element, etree._Element]:
    payloads = _zip_payloads(path)
    document = etree.fromstring(payloads["word/document.xml"])
    relationships = etree.fromstring(payloads["word/_rels/document.xml.rels"])
    return payloads, document, relationships


def _target_row(document: etree._Element) -> tuple[etree._Element | None, list[etree._Element]]:
    for row in document.xpath(".//w:tr", namespaces=NS):
        cells = row.xpath("./w:tc", namespaces=NS)
        if len(cells) >= 2:
            label = "".join(cells[0].xpath(".//w:t/text()", namespaces=NS)).replace(" ", "")
            if "项目介绍材料" in label:
                return row, cells
    return None, []


def _core(artifact: Path) -> tuple[bool, list[str]]:
    errors: list[str] = []
    if not artifact.is_dir():
        return False, ["交付目录不存在"]
    for name in REQUIRED:
        path = artifact / name
        if not path.is_file() or path.stat().st_size == 0:
            errors.append(f"核心交付缺失或为空：{name}")
            continue
        try:
            if name == DOC_NAME:
                _parse_docx(path)
            else:
                deck = Presentation(path)
                if len(deck.slides) == 0:
                    raise ValueError("PPT 无页面")
        except Exception as exc:
            errors.append(f"核心交付无法解析：{name}: {exc}")
    if not errors:
        source = V1.locate_inputs() / SOURCE_DOC_NAME
        if _sha256(artifact / DOC_NAME) == _sha256(source):
            errors.append("仅将未修改的输入 Word 改名交回，未建立要求的链接")
    return not errors, errors


def _criterion(
    key: str,
    score: float,
    evidence: str,
    earned: float = 0,
    total: float = 1,
) -> dict[str, Any]:
    return {
        "criterion_id": key,
        "score": round(max(0.0, min(1.0, score)), 6),
        "evidence": evidence,
        "method": "deterministic",
        "raw": {"earned_units": round(earned, 6), "total_units": round(total, 6)},
    }


def _fatal(errors: list[str], hashes: list[str]) -> dict[str, Any]:
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes or ["0" * 64],
        "criterion_results": [
            _criterion(key, 0, "；".join(errors), 0, 0)
            for key in (*POSITIVE_WEIGHTS, *PENALTY_WEIGHTS)
        ],
        "reward": 0.0,
        "passed": False,
        "hurdle_failures": ["R001"],
        "penalties": [],
    }


def _ratio(findings: list[tuple[bool, str]]) -> tuple[float, list[str], int, int]:
    earned = sum(bool(ok) for ok, _ in findings)
    return earned / len(findings), [label for ok, label in findings if not ok], earned, len(findings)


def _slide_texts(deck: Presentation) -> list[str]:
    result: list[str] = []
    for slide in deck.slides:
        chunks: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                chunks.append(shape.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    chunks.extend(cell.text for cell in row.cells)
        result.append(re.sub(r"\s+", "", "".join(chunks)))
    return result


def _ppt_fidelity(artifact: Path) -> tuple[float, list[str], int, int]:
    inputs = V1.locate_inputs()
    candidate_path = artifact / PPT_NAME
    source_path = inputs / PPT_NAME
    candidate = Presentation(candidate_path)
    source = Presentation(source_path)
    candidate_payloads = _zip_payloads(candidate_path)
    source_payloads = _zip_payloads(source_path)
    candidate_text = _slide_texts(candidate)
    source_text = _slide_texts(source)
    findings: list[tuple[bool, str]] = [
        (_sha256(source_path) == SOURCE_HASHES[PPT_NAME], "冻结 PPT 输入哈希异常"),
        (_sha256(candidate_path) == SOURCE_HASHES[PPT_NAME], "随附 PPT 未逐字节保真"),
        (len(candidate.slides) == 18, f"随附 PPT 页数应为18，实际{len(candidate.slides)}"),
        (set(candidate_payloads) == set(source_payloads), "随附 PPT 包成员集合改变"),
    ]
    for index, expected in enumerate(source_text):
        observed = candidate_text[index] if index < len(candidate_text) else ""
        findings.append((observed == expected, f"第{index + 1}页文字或表格内容改变"))
    source_media = sorted(name for name in source_payloads if name.startswith("ppt/media/"))
    findings.extend(
        (
            candidate_payloads.get(name) == source_payloads[name],
            f"媒体部件 {name} 改变或缺失",
        )
        for name in source_media
    )
    return _ratio(findings)


def _link_context(artifact: Path) -> dict[str, Any]:
    payloads, document, relationships = _parse_docx(artifact / DOC_NAME)
    row, cells = _target_row(document)
    right = cells[1] if len(cells) >= 2 else None
    left = cells[0] if cells else None
    hyperlinks = right.xpath(".//w:hyperlink", namespaces=NS) if right is not None else []
    all_hyperlinks = document.xpath(".//w:hyperlink", namespaces=NS)
    display = "".join(hyperlinks[0].xpath(".//w:t/text()", namespaces=NS)) if len(hyperlinks) == 1 else ""
    relation_id = hyperlinks[0].get(f"{{{R_NS}}}id") if len(hyperlinks) == 1 else None
    found = relationships.xpath("./rel:Relationship[@Id=$rid]", namespaces=NS, rid=relation_id) if relation_id else []
    relation = found[0] if len(found) == 1 else None
    return {
        "payloads": payloads,
        "document": document,
        "relationships": relationships,
        "row": row,
        "cells": cells,
        "left": left,
        "right": right,
        "hyperlinks": hyperlinks,
        "all_hyperlinks": all_hyperlinks,
        "display": display,
        "relation_id": relation_id,
        "relation_matches": found,
        "relation": relation,
    }


def _location_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    ctx = _link_context(artifact)
    left_text = "".join(ctx["left"].xpath(".//w:t/text()", namespaces=NS)).replace(" ", "") if ctx["left"] is not None else ""
    display = ctx["display"]
    findings = [
        (ctx["row"] is not None, "未找到‘项目介绍材料’目标行"),
        (len(ctx["cells"]) == 2, f"目标行应有2个单元格，实际{len(ctx['cells'])}"),
        (left_text == "项目介绍材料", f"左侧标签改变：{left_text!r}"),
        (len(ctx["all_hyperlinks"]) == 1, f"整份 Word 应恰有1个超链接，实际{len(ctx['all_hyperlinks'])}"),
        (len(ctx["hyperlinks"]) == 1, f"右侧单元格应恰有1个超链接，实际{len(ctx['hyperlinks'])}"),
        ("项目介绍材料PPT" in display, "显示文字缺少‘项目介绍材料PPT’"),
        ("青梧谷灵MR共创计划" in display, "显示文字缺少真实项目名"),
        (display.endswith(".pptx"), "显示文字未以 .pptx 结尾"),
    ]
    return _ratio(findings)


def _relationship_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    ctx = _link_context(artifact)
    relation = ctx["relation"]
    target = relation.get("Target", "") if relation is not None else ""
    mode = relation.get("TargetMode", "") if relation is not None else ""
    rel_type = relation.get("Type", "") if relation is not None else ""
    pure = PurePosixPath(target) if target else None
    safe = bool(target) and not pure.is_absolute() and ".." not in pure.parts and ":" not in target and "\\" not in target
    findings = [
        (bool(ctx["relation_id"]), "超链接缺少关系 ID"),
        (len(ctx["relation_matches"]) == 1, f"关系 ID 应唯一匹配，实际{len(ctx['relation_matches'])}"),
        (rel_type.endswith("/hyperlink"), f"关系类型不是 hyperlink：{rel_type!r}"),
        (target == PPT_NAME, f"关系 Target 应为 {PPT_NAME!r}，实际{target!r}"),
        (mode == "External", f"TargetMode 应为 External，实际{mode!r}"),
        (safe, f"关系目标不是安全同目录相对路径：{target!r}"),
    ]
    return _ratio(findings)


def _style_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    ctx = _link_context(artifact)
    hyperlink = ctx["hyperlinks"][0] if len(ctx["hyperlinks"]) == 1 else None
    right = ctx["right"]
    run_properties = hyperlink.xpath(".//w:rPr", namespaces=NS) if hyperlink is not None else []
    first = run_properties[0] if run_properties else None
    underline = first.xpath("./w:u/@w:val", namespaces=NS) if first is not None else []
    sizes = first.xpath("./w:sz/@w:val", namespaces=NS) if first is not None else []
    colors = first.xpath("./w:color/@w:val", namespaces=NS) if first is not None else []
    paragraph_alignment = right.xpath("./w:p/w:pPr/w:jc/@w:val", namespaces=NS) if right is not None else []
    drawings = right.xpath(".//w:drawing|.//w:object|.//w:pict", namespaces=NS) if right is not None else []
    history = hyperlink.get(f"{{{W_NS}}}history") if hyperlink is not None else None
    findings = [
        (hyperlink is not None, "不是可编辑的 w:hyperlink"),
        (not drawings, f"右侧单元格含{len(drawings)}个静态图/对象伪装"),
        (underline in (["single"], ["1"]), f"下划线应为 single，实际{underline}"),
        (sizes == ["18"], f"字号半点值应为18，实际{sizes}"),
        (colors == ["0563C1"], f"链接颜色应为0563C1，实际{colors}"),
        (paragraph_alignment == ["center"], f"链接段落应居中，实际{paragraph_alignment}"),
        (history == "1", f"超链接 history 应为1，实际{history!r}"),
    ]
    return _ratio(findings)


def _preservation_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    inputs = V1.locate_inputs()
    source_payloads, source_document, _ = _parse_docx(inputs / SOURCE_DOC_NAME)
    output_payloads, output_document, _ = _parse_docx(artifact / DOC_NAME)
    allowed = {"word/document.xml", "word/_rels/document.xml.rels"}
    source_text = "".join(source_document.xpath(".//w:t/text()", namespaces=NS))
    output_text = "".join(output_document.xpath(".//w:t/text()", namespaces=NS))
    ctx = _link_context(artifact)
    findings: list[tuple[bool, str]] = [
        (_sha256(inputs / SOURCE_DOC_NAME) == SOURCE_HASHES[SOURCE_DOC_NAME], "冻结 Word 输入哈希异常"),
        (set(source_payloads) <= set(output_payloads), "原 Word 包成员被删除"),
        (source_text in output_text, "原 Word 文字未完整保留"),
        (output_text == source_text + ctx["display"], "Word 出现非要求的新增、删除或改写文字"),
    ]
    for name in sorted(set(source_payloads) - allowed):
        findings.append((output_payloads.get(name) == source_payloads[name], f"非目标部件被改动：{name}"))
    return _ratio(findings)


def _render_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    with tempfile.TemporaryDirectory(prefix="officeval059-render-") as temporary:
        root = Path(temporary)
        output = root / "pdf"
        profile = root / "lo-profile"
        output.mkdir()
        command = [
            "soffice", "--headless", f"-env:UserInstallation={profile.as_uri()}",
            "--convert-to", "pdf", "--outdir", str(output), str(artifact / DOC_NAME),
        ]
        completed = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=120)
        pdf = output / f"{Path(DOC_NAME).stem}.pdf"
        render_ok = completed.returncode == 0 and pdf.is_file() and pdf.stat().st_size > 0
        pages = 0
        page_text = ""
        if render_ok:
            info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=30)
            match = re.search(r"^Pages:\s+(\d+)", info.stdout, re.MULTILINE)
            pages = int(match.group(1)) if match else 0
            text_path = root / "page4.txt"
            extract = subprocess.run(
                ["pdftotext", "-f", "4", "-l", "4", str(pdf), str(text_path)],
                stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=30,
            )
            if extract.returncode == 0 and text_path.is_file():
                page_text = re.sub(r"\s+", "", text_path.read_text(errors="replace"))
        findings = [
            (render_ok, f"Word 转 PDF 失败：{completed.stdout[-200:]}"),
            (pages == 4, f"渲染页数应为4，实际{pages}"),
            ("项目介绍材料" in page_text, "第4页缺少项目介绍材料标签"),
            ("青梧谷灵MR共创计划.pptx" in page_text, "第4页缺少完整链接文字"),
            ("项目申报书" in page_text, "第4页缺少申报书页眉/标题"),
            ("第4页" in page_text, "第4页缺少页码"),
        ]
        return _ratio(findings)


def _table_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    ctx = _link_context(artifact)
    document = ctx["document"]
    tables = document.xpath(".//w:tbl", namespaces=NS)
    application_rows = document.xpath(".//w:tbl[2]/w:tr", namespaces=NS)
    cells = ctx["cells"]
    bottom_borders = [
        bool(cell.xpath("./w:tcPr/w:tcBorders/w:bottom", namespaces=NS))
        for cell in cells[:2]
    ]
    comments = [name for name in ctx["payloads"] if name.startswith("word/comments")]
    objects = ctx["right"].xpath(".//w:drawing|.//w:object|.//w:pict", namespaces=NS) if ctx["right"] is not None else []
    findings = [
        (len(tables) >= 2, f"Word 至少应有2张表，实际{len(tables)}"),
        (len(application_rows) == 25, f"申报表应有25行，实际{len(application_rows)}"),
        (len(cells) == 2, f"目标行应保持2个单元格，实际{len(cells)}"),
        (bottom_borders == [True, True], f"目标行两个单元格底边应完整，实际{bottom_borders}"),
        (not comments, f"Word 含批注部件：{comments}"),
        (not objects, f"目标单元格含静态对象：{len(objects)}"),
    ]
    return _ratio(findings)


def _security_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    actual_files = {path.name for path in artifact.iterdir() if path.is_file()}
    directories = [path.name for path in artifact.iterdir() if path.is_dir()]
    doc_payloads = _zip_payloads(artifact / DOC_NAME)
    ppt_payloads = _zip_payloads(artifact / PPT_NAME)
    ctx = _link_context(artifact)
    external = [
        rel for rel in ctx["relationships"].xpath("./rel:Relationship", namespaces=NS)
        if rel.get("TargetMode") == "External"
    ]
    macro_parts = [
        name for name in (*doc_payloads, *ppt_payloads)
        if name.endswith("vbaProject.bin") or name.endswith("vbaData.xml")
    ]
    target = ctx["relation"].get("Target", "") if ctx["relation"] is not None else ""
    findings = [
        (actual_files == set(REQUIRED) and not directories, f"交付文件或子目录不符合要求：文件={sorted(actual_files)}，目录={directories}"),
        (not macro_parts, f"OOXML 含宏部件：{macro_parts}"),
        (len(external) == 1, f"Word 应恰有1条 External 关系，实际{len(external)}"),
        (all(rel.get("Target", "") == PPT_NAME for rel in external), "存在非指定外部关系"),
        (not re.match(r"^(?:[A-Za-z]:|/|\\\\)", target), f"链接泄露绝对路径：{target!r}"),
        (not any(name.startswith("word/comments") for name in doc_payloads), "Word 含批注"),
    ]
    return _ratio(findings)


def _unsafe_target_penalty(artifact: Path) -> tuple[float, list[str], float]:
    ctx = _link_context(artifact)
    relation = ctx["relation"]
    target = relation.get("Target", "") if relation is not None else ""
    if not target:
        return 0.45, ["目标关系缺失，链接不可点击"], 1.0
    pure = PurePosixPath(target)
    unsafe = pure.is_absolute() or ".." in pure.parts or ":" in target or "\\" in target
    if unsafe:
        return 1.0, [f"链接目标越界或泄露绝对路径：{target!r}"], 1.0
    if target != PPT_NAME:
        return 0.65, [f"链接指向非指定文件：{target!r}"], 1.0
    if relation.get("TargetMode") != "External" or not relation.get("Type", "").endswith("/hyperlink"):
        return 0.4, ["链接关系类型或 TargetMode 错误"], 1.0
    return 0.0, [], 1.0


def _extra_penalty(artifact: Path) -> tuple[float, list[str], float]:
    actual = {path.name for path in artifact.iterdir() if path.is_file()}
    directories = [path.name for path in artifact.iterdir() if path.is_dir()]
    extras = sorted(actual - set(REQUIRED)) + directories
    doc_payloads = _zip_payloads(artifact / DOC_NAME)
    ppt_payloads = _zip_payloads(artifact / PPT_NAME)
    macro = any(name.endswith("vbaProject.bin") or name.endswith("vbaData.xml") for name in (*doc_payloads, *ppt_payloads))
    comments = any(name.startswith("word/comments") for name in doc_payloads)
    labels: list[str] = []
    if extras:
        labels.append(f"额外交付或子目录：{extras}")
    if macro:
        labels.append("OOXML 含宏")
    if comments:
        labels.append("Word 含批注")
    severity = min(1.0, 0.5 * bool(extras) + 0.75 * bool(macro) + 0.5 * bool(comments))
    return severity, labels, 1.0


def _safe_positive(
    call: Callable[[Path], tuple[float, list[str], int, int]],
    artifact: Path,
) -> tuple[float, list[str], float, float]:
    try:
        return call(artifact)
    except Exception as exc:
        return 0.0, [f"任务专属检查无法完成：{type(exc).__name__}: {exc}"], 0, 1


def _safe_penalty(
    call: Callable[[Path], tuple[float, list[str], float]],
    artifact: Path,
) -> tuple[float, list[str], float]:
    try:
        return call(artifact)
    except Exception as exc:
        return 1.0, [f"严重边界无法核验：{type(exc).__name__}: {exc}"], 1.0


def evaluate(artifact: Path) -> dict[str, Any]:
    ok, errors = _core(artifact)
    hashes = [_sha256(artifact / name) for name in REQUIRED if (artifact / name).is_file()]
    if not ok:
        return _fatal(errors, hashes)

    custom = {
        "R002": _ppt_fidelity,
        "R003": _location_positive,
        "R004": _relationship_positive,
        "R005": _style_positive,
        "R006": _preservation_positive,
        "R007": _render_positive,
        "R008": _table_positive,
        "R009": _security_positive,
    }
    results = [_criterion("R001", 1, "指定 DOCX/PPTX 存在、非空且可解析，未原样交回输入 Word", 4, 4)]
    scores: dict[str, float] = {"R001": 1.0}
    for key in [f"R{index:03d}" for index in range(2, 10)]:
        score, failures, earned, total = _safe_positive(custom[key], artifact)
        results.append(_criterion(
            key,
            score,
            "本题任务专属检查通过" if not failures else "未通过：" + "；".join(failures),
            earned,
            total,
        ))
        scores[key] = score

    for key, call in (("P001", _unsafe_target_penalty), ("P002", _extra_penalty)):
        score, labels, total = _safe_penalty(call, artifact)
        results.append(_criterion(
            key,
            score,
            "未触发任务专属严重边界" if not labels else "；".join(labels),
            score * total,
            total,
        ))
        scores[key] = score

    positive_total = sum(POSITIVE_WEIGHTS.values())
    positive_earned = sum(POSITIVE_WEIGHTS[key] * scores[key] for key in POSITIVE_WEIGHTS)
    penalty_points = sum(abs(PENALTY_WEIGHTS[key]) * scores[key] for key in PENALTY_WEIGHTS)
    reward = round(max(0.0, min(1.0, (positive_earned - penalty_points) / positive_total)), 6)
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes,
        "criterion_results": results,
        "reward": reward,
        "passed": reward >= 0.8,
        "hurdle_failures": [],
        "penalties": [
            {"id": key, "amount": round(abs(PENALTY_WEIGHTS[key]) * scores[key] / positive_total, 6)}
            for key in PENALTY_WEIGHTS if scores[key] > 0
        ],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    parser.add_argument("--scoring-version", choices=("v1", "v2"), default="v2")
    args = parser.parse_args()
    try:
        result = V1.evaluate(args.artifact_dir) if args.scoring_version == "v1" else evaluate(args.artifact_dir)
    except Exception as exc:
        result = _fatal([f"Judge 异常：{type(exc).__name__}: {exc}"], [])
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{float(result['reward']):.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
