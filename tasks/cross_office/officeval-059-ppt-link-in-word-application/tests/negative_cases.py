#!/usr/bin/env python3
"""OfficeVal 059 的任务专属 v2 分层负例。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Callable

from lxml import etree
from docx import Document
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
INPUTS = ROOT / "environment" / "input_files"
DOC_NAME = "青梧谷灵MR共创计划_项目申报书_已链接PPT.docx"
SOURCE_DOC_NAME = "青梧谷灵MR共创计划_项目申报书.docx"
PPT_NAME = "青梧谷灵MR共创计划.pptx"
OUTPUTS = (DOC_NAME, PPT_NAME)
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"w": W_NS, "r": R_NS, "rel": REL_NS}
Mutator = Callable[[Path], None]


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"无法载入 {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _rewrite_zip(path: Path, replacements: dict[str, bytes]) -> None:
    with zipfile.ZipFile(path) as source:
        entries = [(item, source.read(item.filename)) for item in source.infolist()]
    temporary = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(temporary, "w") as target:
        for item, payload in entries:
            target.writestr(item, replacements.get(item.filename, payload))
    temporary.replace(path)


def _doc_roots(root: Path) -> tuple[etree._Element, etree._Element]:
    with zipfile.ZipFile(root / DOC_NAME) as archive:
        document = etree.fromstring(archive.read("word/document.xml"))
        relationships = etree.fromstring(archive.read("word/_rels/document.xml.rels"))
    return document, relationships


def _target(root: Path) -> tuple[etree._Element, list[etree._Element], etree._Element, etree._Element]:
    document, relationships = _doc_roots(root)
    for row in document.xpath(".//w:tr", namespaces=NS):
        cells = row.xpath("./w:tc", namespaces=NS)
        if len(cells) >= 2 and "项目介绍材料" in "".join(cells[0].xpath(".//w:t/text()", namespaces=NS)):
            hyperlinks = cells[1].xpath(".//w:hyperlink", namespaces=NS)
            if len(hyperlinks) != 1:
                raise ValueError("Golden 目标单元格没有唯一超链接")
            relation_id = hyperlinks[0].get(f"{{{R_NS}}}id")
            found = relationships.xpath("./rel:Relationship[@Id=$rid]", namespaces=NS, rid=relation_id)
            if len(found) != 1:
                raise ValueError("Golden 超链接关系不唯一")
            return document, cells, hyperlinks[0], found[0]
    raise ValueError("Golden 中找不到项目介绍材料行")


def _save_doc(root: Path, document: etree._Element, relationships: etree._Element) -> None:
    _rewrite_zip(root / DOC_NAME, {
        "word/document.xml": etree.tostring(document, xml_declaration=True, encoding="UTF-8", standalone=True),
        "word/_rels/document.xml.rels": etree.tostring(relationships, xml_declaration=True, encoding="UTF-8", standalone=True),
    })


def _display(root: Path, text: str) -> None:
    document, relationships = _doc_roots(root)
    hyperlinks = document.xpath(".//w:hyperlink", namespaces=NS)
    hyperlinks[0].xpath(".//w:t", namespaces=NS)[0].text = text
    _save_doc(root, document, relationships)


def _left_label(root: Path) -> None:
    document, relationships = _doc_roots(root)
    for node in document.xpath(".//w:t[text()='项目介绍材料']", namespaces=NS):
        node.text = "项目背景材料"
        break
    _save_doc(root, document, relationships)


def _move_link_left(root: Path) -> None:
    document, cells, hyperlink, relationships_item = _target(root)
    _, relationships = _doc_roots(root)
    hyperlink.getparent().remove(hyperlink)
    cells[0].xpath("./w:p", namespaces=NS)[0].append(hyperlink)
    _save_doc(root, document, relationships)


def _relationship(root: Path, attribute: str, value: str | None) -> None:
    document, relationships = _doc_roots(root)
    relation = relationships.xpath("./rel:Relationship[@TargetMode='External']", namespaces=NS)[0]
    if value is None:
        relation.attrib.pop(attribute, None)
    else:
        relation.set(attribute, value)
    _save_doc(root, document, relationships)


def _remove_relationship(root: Path) -> None:
    document, relationships = _doc_roots(root)
    relation = relationships.xpath("./rel:Relationship[@TargetMode='External']", namespaces=NS)[0]
    relationships.remove(relation)
    _save_doc(root, document, relationships)


def _style(root: Path, element: str, attribute: str, value: str | None) -> None:
    document, relationships = _doc_roots(root)
    node = document.xpath(f".//w:hyperlink//w:rPr/w:{element}", namespaces=NS)[0]
    qualified = f"{{{W_NS}}}{attribute}"
    if value is None:
        node.getparent().remove(node)
    else:
        node.set(qualified, value)
    _save_doc(root, document, relationships)


def _alignment(root: Path) -> None:
    document, relationships = _doc_roots(root)
    node = document.xpath(".//w:hyperlink/ancestor::w:p[1]/w:pPr/w:jc", namespaces=NS)[0]
    node.set(f"{{{W_NS}}}val", "left")
    _save_doc(root, document, relationships)


def _history(root: Path) -> None:
    document, relationships = _doc_roots(root)
    hyperlink = document.xpath(".//w:hyperlink", namespaces=NS)[0]
    hyperlink.attrib.pop(f"{{{W_NS}}}history", None)
    _save_doc(root, document, relationships)


def _unwrap_link(root: Path) -> None:
    document, relationships = _doc_roots(root)
    hyperlink = document.xpath(".//w:hyperlink", namespaces=NS)[0]
    parent = hyperlink.getparent()
    index = parent.index(hyperlink)
    for child in list(hyperlink):
        hyperlink.remove(child)
        parent.insert(index, child)
        index += 1
    parent.remove(hyperlink)
    relation = relationships.xpath("./rel:Relationship[@TargetMode='External']", namespaces=NS)[0]
    relationships.remove(relation)
    _save_doc(root, document, relationships)


def _change_unrelated_word_text(root: Path) -> None:
    document, relationships = _doc_roots(root)
    candidates = [node for node in document.xpath(".//w:t", namespaces=NS) if node.text and "项目介绍材料" not in node.text]
    candidates[0].text = f"{candidates[0].text}（错误改写）"
    _save_doc(root, document, relationships)


def _change_core_properties(root: Path) -> None:
    path = root / DOC_NAME
    with zipfile.ZipFile(path) as archive:
        core = etree.fromstring(archive.read("docProps/core.xml"))
    core.append(etree.Comment("不应改动的调试痕迹"))
    _rewrite_zip(path, {"docProps/core.xml": etree.tostring(core, xml_declaration=True, encoding="UTF-8", standalone=True)})


def _remove_bottom_borders(root: Path) -> None:
    document, cells, _, _ = _target(root)
    _, relationships = _doc_roots(root)
    for cell in cells[:2]:
        for node in cell.xpath("./w:tcPr/w:tcBorders/w:bottom", namespaces=NS):
            node.getparent().remove(node)
    _save_doc(root, document, relationships)


def _ppt_title(root: Path) -> None:
    path = root / PPT_NAME
    deck = Presentation(path)
    for shape in deck.slides[0].shapes:
        if getattr(shape, "has_text_frame", False) and "青梧谷灵" in shape.text:
            shape.text = shape.text.replace("青梧谷灵", "错误项目", 1)
            deck.save(path)
            return
    raise ValueError("PPT 首页找不到项目标题")


def _remove_last_slide(root: Path) -> None:
    path = root / PPT_NAME
    deck = Presentation(path)
    slide_id = deck.slides._sldIdLst[-1]
    relation_id = slide_id.rId
    deck.part.drop_rel(relation_id)
    deck.slides._sldIdLst.remove(slide_id)
    deck.save(path)


def _input_only(root: Path) -> None:
    for path in list(root.iterdir()):
        if path.is_file():
            path.unlink()
    shutil.copy2(INPUTS / PPT_NAME, root / PPT_NAME)
    shutil.copy2(INPUTS / SOURCE_DOC_NAME, root / DOC_NAME)


def _severe_incomplete(root: Path) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "青梧谷灵 MR共创计划"
    deck.save(root / PPT_NAME)
    document = Document()
    document.add_heading("青梧谷灵项目申报书", level=1)
    document.add_paragraph("项目介绍材料缺失，未建立可点击链接。")
    document.save(root / DOC_NAME)


def _cases() -> list[tuple[str, Mutator, str]]:
    cases: list[tuple[str, Mutator, str]] = []
    for name in OUTPUTS:
        cases.extend([
            (f"missing-{name}", lambda root, n=name: (root / n).unlink(), "fatal_zero"),
            (f"empty-{name}", lambda root, n=name: (root / n).write_bytes(b""), "fatal_zero"),
            (f"corrupt-{name}", lambda root, n=name: (root / n).write_bytes(b"not-a-valid-office-file"), "fatal_zero"),
        ])
    cases.extend([
        ("empty-output", lambda root: [path.unlink() for path in list(root.iterdir()) if path.is_file()], "fatal_zero"),
        ("unmodified-input-renamed", _input_only, "fatal_zero"),
        ("display-project-name-wrong", lambda root: _display(root, "项目介绍材料PPT：错误项目.pptx"), "partial"),
        ("display-prefix-missing", lambda root: _display(root, "青梧谷灵MR共创计划.pptx"), "partial"),
        ("left-label-changed", _left_label, "partial"),
        ("hyperlink-moved-left", _move_link_left, "partial"),
        ("target-wrong-file", lambda root: _relationship(root, "Target", "其他项目.pptx"), "partial"),
        ("target-absolute-path", lambda root: _relationship(root, "Target", "/tmp/青梧谷灵MR共创计划.pptx"), "partial"),
        ("target-parent-path", lambda root: _relationship(root, "Target", "../青梧谷灵MR共创计划.pptx"), "partial"),
        ("target-mode-missing", lambda root: _relationship(root, "TargetMode", None), "partial"),
        ("relationship-type-wrong", lambda root: _relationship(root, "Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"), "partial"),
        ("underline-removed", lambda root: _style(root, "u", "val", None), "partial"),
        ("font-size-wrong", lambda root: _style(root, "sz", "val", "28"), "partial"),
        ("link-color-wrong", lambda root: _style(root, "color", "val", "FF0000"), "partial"),
        ("paragraph-left-aligned", _alignment, "partial"),
        ("hyperlink-history-removed", _history, "partial"),
        ("ppt-cover-project-wrong", _ppt_title, "partial"),
        ("ppt-last-slide-removed", _remove_last_slide, "partial"),
        ("unrelated-word-text-changed", _change_unrelated_word_text, "partial"),
        ("unrelated-core-properties-changed", _change_core_properties, "partial"),
        ("target-row-bottom-borders-removed", _remove_bottom_borders, "partial"),
        ("extra-debug-file", lambda root: (root / "debug.txt").write_text("debug", encoding="utf-8"), "partial"),
        ("moderate-visible-text-without-link", _unwrap_link, "moderate"),
        ("severe-valid-incomplete-package", _severe_incomplete, "severe"),
    ])
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    judge = _load("officeval_059_v2_judge", ROOT / "tests" / "judge.py")
    if judge.evaluate(GOLDEN).get("reward") != 1.0:
        raise AssertionError("Golden v2 必须为 1.0")
    results: list[dict] = []
    for name, mutator, layer in _cases():
        with tempfile.TemporaryDirectory(prefix=f"officeval-059-v2-{name}-") as temporary:
            output = Path(temporary) / "output"
            shutil.copytree(GOLDEN, output)
            mutator(output)
            result = judge.evaluate(output)
            reward = float(result.get("reward", 0))
            if layer == "fatal_zero" and reward != 0:
                raise AssertionError(f"{name} 应硬归零，实际 {reward}")
            if layer in {"partial", "moderate", "severe"} and not (0 < reward < 1):
                raise AssertionError(f"{name} 应保留连续部分分，实际 {reward}")
            if layer == "moderate" and reward >= 0.8:
                raise AssertionError(f"{name} 应低于通过线，实际 {reward}")
            results.append({"case": name, "layer": layer, "reward": reward, "passed": bool(result.get("passed")), "penalties": result.get("penalties", [])})
    moderate = next(item["reward"] for item in results if item["layer"] == "moderate")
    severe = next(item["reward"] for item in results if item["layer"] == "severe")
    if not severe < moderate:
        raise AssertionError(f"严重负例 {severe} 应低于中度负例 {moderate}")
    partial = [item["reward"] for item in results if item["layer"] != "fatal_zero"]
    report = {
        "total": len(results),
        "fatal_zero": sum(item["layer"] == "fatal_zero" for item in results),
        "partial_continuous": len(partial),
        "partial_reward_range": [min(partial), max(partial)],
        "moderate_reward": moderate,
        "severe_reward": severe,
        "cases": results,
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
