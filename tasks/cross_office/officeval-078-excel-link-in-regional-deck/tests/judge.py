#!/usr/bin/env python3
"""OfficeVal 078 的 v2 连续评分入口；原严格口径保存在 judge_v1.py。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import re
import subprocess
import sys
import tempfile
import warnings
import zipfile
from pathlib import Path, PurePosixPath
from typing import Any, Callable

from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation


TASK_ID = "cowork-evo/officeval-078-excel-link-in-regional-deck"
RUBRIC_VERSION = "omegause-officeval-078-adaptation-v0.2"
JUDGE_VERSION = "officeval-078-criterion-partial-v2.0"
PPT_NAME = "区域协同汇报_第8页Excel超链接版.pptx"
SOURCE_PPT_NAME = "区域协同汇报.pptx"
XLSX_NAME = "CRM_pipeline_restyled_reordered.xlsx"
REQUIRED = (PPT_NAME, XLSX_NAME)
SOURCE_HASHES = {
    SOURCE_PPT_NAME: "3eddccca700d7c3f0c30f1cdb747cbf984a68757583a81aef524edc87a99c42b",
    XLSX_NAME: "e45c7f9aac5c65df4e28e67d503a9aac4712b31b59ff219b234fa85d388a1839",
}
EXPECTED_SHEETS = {
    "Strategic Pipeline": (160, 7, "StrategicPipelineTbl"),
    "Lead Development": (220, 24, "LeadDevelopmentTbl"),
    "Client Portfolio": (480, 43, "ClientPortfolioTbl"),
}
POSITIVE_WEIGHTS = {
    "R001": 6,
    "R002": 13,
    "R003": 14,
    "R004": 12,
    "R005": 12,
    "R006": 14,
    "R007": 12,
    "R008": 9,
    "R009": 8,
}
PENALTY_WEIGHTS = {"P001": -12, "P002": -4}
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"p": P_NS, "a": A_NS, "r": R_NS, "rel": REL_NS}


def _load_v1():
    path = Path(__file__).with_name("judge_v1.py")
    spec = importlib.util.spec_from_file_location("officeval_078_judge_v1", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v1 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


V1 = _load_v1()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _zip_payloads(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as archive:
        bad = archive.testzip()
        if bad:
            raise ValueError(f"OOXML 部件损坏：{bad}")
        return {name: archive.read(name) for name in archive.namelist()}


def _open_workbook(path: Path):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return load_workbook(path, data_only=False, read_only=False)


def _core(artifact: Path) -> tuple[bool, list[str]]:
    errors: list[str] = []
    if not artifact.is_dir():
        return False, ["交付目录不存在"]
    for name in REQUIRED:
        path = artifact / name
        if not path.is_file() or path.stat().st_size == 0:
            errors.append(f"核心交付缺失或为空：{name}")
            continue
        try:
            if name == PPT_NAME:
                deck = Presentation(path)
                if len(deck.slides) == 0:
                    raise ValueError("PPT 无页面")
                _zip_payloads(path)
            else:
                book = _open_workbook(path)
                if not book.worksheets:
                    raise ValueError("工作簿无工作表")
                book.close()
        except Exception as exc:
            errors.append(f"核心交付无法解析：{name}: {exc}")
    if not errors:
        inputs = V1.locate_inputs()
        candidate_payloads = _zip_payloads(artifact / PPT_NAME)
        source_payloads = _zip_payloads(inputs / SOURCE_PPT_NAME)
        if candidate_payloads == source_payloads:
            errors.append("仅将未修改的输入 PPT 改名交回，未建立要求的 Excel 链接按钮")
    return not errors, errors


def _criterion(key: str, score: float, evidence: str, earned: float = 0, total: float = 1) -> dict[str, Any]:
    return {
        "criterion_id": key,
        "score": round(max(0.0, min(1.0, score)), 6),
        "evidence": evidence,
        "method": "deterministic",
        "raw": {"earned_units": round(earned, 6), "total_units": round(total, 6)},
    }


def _fatal(errors: list[str], hashes: list[str]) -> dict[str, Any]:
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes or ["0" * 64],
        "criterion_results": [_criterion(key, 0, "；".join(errors), 0, 0) for key in (*POSITIVE_WEIGHTS, *PENALTY_WEIGHTS)],
        "reward": 0.0,
        "passed": False,
        "hurdle_failures": ["R001"],
        "penalties": [],
    }


def _ratio(findings: list[tuple[bool, str]]) -> tuple[float, list[str], int, int]:
    earned = sum(bool(ok) for ok, _ in findings)
    return earned / len(findings), [label for ok, label in findings if not ok], earned, len(findings)


def _sheet_values(sheet) -> list[list[Any]]:
    return [[cell.value for cell in row] for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=sheet.max_column)]


def _workbook_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    inputs = V1.locate_inputs()
    source_path = inputs / XLSX_NAME
    candidate_path = artifact / XLSX_NAME
    source = _open_workbook(source_path)
    candidate = _open_workbook(candidate_path)
    findings: list[tuple[bool, str]] = [
        (_sha256(source_path) == SOURCE_HASHES[XLSX_NAME], "冻结工作簿输入哈希异常"),
        (_sha256(candidate_path) == SOURCE_HASHES[XLSX_NAME], "随附工作簿未逐字节保真"),
        (candidate.sheetnames == list(EXPECTED_SHEETS), f"工作表名称/顺序错误：{candidate.sheetnames}"),
    ]
    for name, (rows, columns, table_name) in EXPECTED_SHEETS.items():
        observed = candidate[name] if name in candidate.sheetnames else None
        expected = source[name]
        findings.extend([
            (observed is not None and observed.max_row == rows, f"{name} 行数应为{rows}"),
            (observed is not None and observed.max_column == columns, f"{name} 列数应为{columns}"),
            (observed is not None and table_name in observed.tables, f"{name} 缺少原生表 {table_name}"),
            (observed is not None and [observed.cell(3, column).value for column in range(1, columns + 1)] == [expected.cell(3, column).value for column in range(1, columns + 1)], f"{name} 表头改变"),
            (observed is not None and _sheet_values(observed) == _sheet_values(expected), f"{name} 记录或公式改变"),
        ])
    source.close()
    candidate.close()
    return _ratio(findings)


def _link_context(artifact: Path) -> dict[str, Any]:
    payloads = _zip_payloads(artifact / PPT_NAME)
    slide = etree.fromstring(payloads["ppt/slides/slide8.xml"])
    relationships = etree.fromstring(payloads["ppt/slides/_rels/slide8.xml.rels"])
    shape_links = slide.xpath(".//p:cNvPr/a:hlinkClick", namespaces=NS)
    run_links = slide.xpath(".//a:rPr/a:hlinkClick", namespaces=NS)
    all_links: list[tuple[int, etree._Element]] = []
    for index in range(1, 11):
        member = f"ppt/slides/slide{index}.xml"
        if member in payloads:
            root = etree.fromstring(payloads[member])
            all_links.extend((index, item) for item in root.xpath(".//a:hlinkClick", namespaces=NS))
    link = shape_links[0] if len(shape_links) == 1 else None
    shapes = link.xpath("ancestor::p:sp[1]", namespaces=NS) if link is not None else []
    shape = shapes[0] if len(shapes) == 1 else None
    relation_id = link.get(f"{{{R_NS}}}id") if link is not None else None
    matches = relationships.xpath("./rel:Relationship[@Id=$rid]", namespaces=NS, rid=relation_id) if relation_id else []
    relation = matches[0] if len(matches) == 1 else None
    text = "".join(shape.xpath(".//a:t/text()", namespaces=NS)) if shape is not None else ""
    return {
        "payloads": payloads,
        "slide": slide,
        "relationships": relationships,
        "shape_links": shape_links,
        "run_links": run_links,
        "all_links": all_links,
        "link": link,
        "shape": shape,
        "relation_id": relation_id,
        "relation_matches": matches,
        "relation": relation,
        "text": text,
    }


def _button_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    ctx = _link_context(artifact)
    shape = ctx["shape"]
    text = re.sub(r"\s+", "", ctx["text"])
    name = shape.xpath("./p:nvSpPr/p:cNvPr/@name", namespaces=NS) if shape is not None else []
    link_slides = [index for index, _ in ctx["all_links"]]
    findings = [
        (len(ctx["all_links"]) == 1, f"整份 PPT 应恰有1个点击链接，实际{len(ctx['all_links'])}"),
        (link_slides == [8], f"点击链接应仅位于第8页，实际{link_slides}"),
        (len(ctx["shape_links"]) == 1, f"第8页形状级链接应为1，实际{len(ctx['shape_links'])}"),
        (not ctx["run_links"], f"不得用文字 run 级链接替代形状级链接，实际{len(ctx['run_links'])}"),
        (shape is not None, "形状级链接未绑定到可编辑形状"),
        ("打开" in text and "Excel" in text, f"按钮文字未明确表达打开 Excel：{ctx['text']!r}"),
        ("CRM" in text and "工作簿" in text, f"按钮文字未识别 CRM 工作簿：{ctx['text']!r}"),
        (bool(name) and "CRM_pipeline_restyled_reordered.xlsx" in name[0], f"按钮可访问名称未记录目标文件：{name}"),
    ]
    return _ratio(findings)


def _geometry_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    ctx = _link_context(artifact)
    shape = ctx["shape"]
    xfrm = shape.find("./p:spPr/a:xfrm", namespaces=NS) if shape is not None else None
    offset = xfrm.find("a:off", namespaces=NS) if xfrm is not None else None
    extent = xfrm.find("a:ext", namespaces=NS) if xfrm is not None else None
    if offset is not None and extent is not None:
        x = int(offset.get("x")) / 360000
        y = int(offset.get("y")) / 360000
        width = int(extent.get("cx")) / 360000
        height = int(extent.get("cy")) / 360000
    else:
        x = y = width = height = 0.0
    findings = [
        (26.0 <= x <= 29.0, f"按钮左边缘应为26–29cm，实际{x:.2f}"),
        (5.0 <= y <= 11.0, f"按钮上边缘应为5–11cm，实际{y:.2f}"),
        (3.5 <= width <= 5.5, f"按钮宽度应为3.5–5.5cm，实际{width:.2f}"),
        (1.5 <= height <= 3.5, f"按钮高度应为1.5–3.5cm，实际{height:.2f}"),
        (x >= 25.65, f"按钮侵入左侧表格区：x={x:.2f}"),
        (x + width <= 32.17, f"按钮右缘过近或超出95%边界：{x + width:.2f}"),
        (y >= 3.5, f"按钮进入顶部标题区：y={y:.2f}"),
        (y + height < 14.73, f"按钮进入底部说明、页脚或橙色圆形区：{y + height:.2f}"),
        (width * height <= 64.52, f"按钮占用空白区过大：{width * height:.2f}cm²"),
    ]
    return _ratio(findings)


def _relationship_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    ctx = _link_context(artifact)
    relation = ctx["relation"]
    target = relation.get("Target", "") if relation is not None else ""
    mode = relation.get("TargetMode", "") if relation is not None else ""
    rel_type = relation.get("Type", "") if relation is not None else ""
    pure = PurePosixPath(target) if target else None
    safe = bool(target) and not pure.is_absolute() and ".." not in pure.parts and ":" not in target and "\\" not in target
    findings = [
        (bool(ctx["relation_id"]), "按钮链接缺少关系 ID"),
        (len(ctx["relation_matches"]) == 1, f"关系 ID 应唯一匹配，实际{len(ctx['relation_matches'])}"),
        (rel_type.endswith("/hyperlink"), f"关系类型不是 hyperlink：{rel_type!r}"),
        (target == XLSX_NAME, f"Target 应为 {XLSX_NAME!r}，实际{target!r}"),
        (mode == "External", f"TargetMode 应为 External，实际{mode!r}"),
        (safe, f"Target 不是安全同目录相对路径：{target!r}"),
        ((artifact / target).is_file() if safe else False, "相对目标文件未与 PPT 一同交付"),
    ]
    return _ratio(findings)


def _shape_signature(shape) -> tuple[Any, ...]:
    text = shape.text if getattr(shape, "has_text_frame", False) else ""
    table_text: list[str] = []
    if getattr(shape, "has_table", False):
        table_text = [cell.text for row in shape.table.rows for cell in row.cells]
    return (int(shape.shape_type), shape.left, shape.top, shape.width, shape.height, text, tuple(table_text))


def _preservation_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    inputs = V1.locate_inputs()
    source_path = inputs / SOURCE_PPT_NAME
    candidate_path = artifact / PPT_NAME
    source_payloads = _zip_payloads(source_path)
    candidate_payloads = _zip_payloads(candidate_path)
    source = Presentation(source_path)
    candidate = Presentation(candidate_path)
    allowed = {"ppt/slides/slide8.xml", "ppt/slides/_rels/slide8.xml.rels"}
    findings: list[tuple[bool, str]] = [
        (_sha256(source_path) == SOURCE_HASHES[SOURCE_PPT_NAME], "冻结 PPT 输入哈希异常"),
        (set(source_payloads) <= set(candidate_payloads), "原 PPT 包成员被删除"),
        (len(candidate.slides) == 10, f"PPT 应保持10页，实际{len(candidate.slides)}"),
        (candidate.slide_width == source.slide_width and candidate.slide_height == source.slide_height, "宽屏页面尺寸改变"),
    ]
    for name in sorted(set(source_payloads) - allowed):
        findings.append((candidate_payloads.get(name) == source_payloads[name], f"非目标部件被改动：{name}"))
    for index, source_slide in enumerate(source.slides):
        if index >= len(candidate.slides):
            findings.append((False, f"第{index + 1}页缺失"))
            continue
        source_shapes = [_shape_signature(shape) for shape in source_slide.shapes]
        candidate_shapes = [_shape_signature(shape) for shape in list(candidate.slides[index].shapes)[:len(source_shapes)]]
        findings.append((candidate_shapes == source_shapes, f"第{index + 1}页原对象、文字或几何改变"))
    return _ratio(findings)


def _render_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    with tempfile.TemporaryDirectory(prefix="officeval078-render-") as temporary:
        root = Path(temporary)
        output = root / "pdf"
        profile = root / "lo-profile"
        output.mkdir()
        completed = subprocess.run(
            ["soffice", "--headless", f"-env:UserInstallation={profile.as_uri()}", "--convert-to", "pdf", "--outdir", str(output), str(artifact / PPT_NAME)],
            stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=150,
        )
        pdf = output / f"{Path(PPT_NAME).stem}.pdf"
        render_ok = completed.returncode == 0 and pdf.is_file() and pdf.stat().st_size > 0
        pages = 0
        page_text = ""
        if render_ok:
            info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=30)
            match = re.search(r"^Pages:\s+(\d+)", info.stdout, re.MULTILINE)
            pages = int(match.group(1)) if match else 0
            text_path = root / "slide8.txt"
            extract = subprocess.run(["pdftotext", "-f", "8", "-l", "8", str(pdf), str(text_path)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=30)
            if extract.returncode == 0 and text_path.is_file():
                page_text = re.sub(r"\s+", "", text_path.read_text(errors="replace"))
        required_texts = ["StrategicPipelineStatusUpdate", "HarborLink", "NorthBay", "Pipelineview", "打开Excel", "CRM工作簿", "08/10"]
        findings = [(render_ok, f"PPT 转 PDF 失败：{completed.stdout[-200:] }"), (pages == 10, f"渲染页数应为10，实际{pages}")]
        findings.extend((text in page_text, f"第8页渲染缺少 {text}") for text in required_texts)
        return _ratio(findings)


def _style_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    ctx = _link_context(artifact)
    shape = ctx["shape"]
    geometry = shape.xpath("./p:spPr/a:prstGeom/@prst", namespaces=NS) if shape is not None else []
    fills = shape.xpath("./p:spPr/a:solidFill/a:srgbClr/@val", namespaces=NS) if shape is not None else []
    line_fills = shape.xpath("./p:spPr/a:ln/a:solidFill/a:srgbClr/@val", namespaces=NS) if shape is not None else []
    colors = shape.xpath(".//a:rPr/a:solidFill/a:srgbClr/@val", namespaces=NS) if shape is not None else []
    alignments = shape.xpath(".//a:pPr/@algn", namespaces=NS) if shape is not None else []
    anchors = shape.xpath("./p:txBody/a:bodyPr/@anchor", namespaces=NS) if shape is not None else []
    findings = [
        (shape is not None, "按钮不是可编辑 p:sp 形状"),
        (geometry == ["roundRect"], f"按钮应为 roundRect，实际{geometry}"),
        (fills == ["107C41"], f"按钮填充应为 Excel 绿107C41，实际{fills}"),
        (line_fills == ["33C481"], f"按钮边线应为33C481，实际{line_fills}"),
        (len(colors) >= 2 and all(color == "FFFFFF" for color in colors), f"按钮文字应为白色，实际{colors}"),
        (len(alignments) >= 2 and all(item == "ctr" for item in alignments), f"按钮文字应居中，实际{alignments}"),
        (anchors == ["ctr"], f"按钮垂直锚点应居中，实际{anchors}"),
        (len(ctx["shape_links"]) == 1 and not ctx["run_links"], "整个形状没有完整点击区"),
    ]
    return _ratio(findings)


def _security_positive(artifact: Path) -> tuple[float, list[str], int, int]:
    actual_files = {path.name for path in artifact.iterdir() if path.is_file()}
    directories = [path.name for path in artifact.iterdir() if path.is_dir()]
    payloads = _zip_payloads(artifact / PPT_NAME)
    ctx = _link_context(artifact)
    external = [rel for rel in ctx["relationships"].xpath("./rel:Relationship", namespaces=NS) if rel.get("TargetMode") == "External"]
    macro_parts = [name for name in payloads if name.endswith("vbaProject.bin") or name.endswith("vbaData.xml")]
    comments = [name for name in payloads if "comment" in name.casefold()]
    target = ctx["relation"].get("Target", "") if ctx["relation"] is not None else ""
    findings = [
        (actual_files == set(REQUIRED) and not directories, f"交付文件或子目录错误：文件={sorted(actual_files)}，目录={directories}"),
        (not macro_parts, f"PPT 含宏部件：{macro_parts}"),
        (not comments, f"PPT 含批注：{comments}"),
        (len(external) == 1, f"第8页应恰有1条 External 关系，实际{len(external)}"),
        (all(rel.get("Target", "") == XLSX_NAME for rel in external), "存在非指定外部关系"),
        (not re.match(r"^(?:[A-Za-z]:|/|\\\\)", target), f"链接泄露绝对路径：{target!r}"),
    ]
    return _ratio(findings)


def _unsafe_target_penalty(artifact: Path) -> tuple[float, list[str], float]:
    ctx = _link_context(artifact)
    relation = ctx["relation"]
    target = relation.get("Target", "") if relation is not None else ""
    if not target:
        return 0.75, ["目标关系缺失，按钮不可点击"], 1.0
    pure = PurePosixPath(target)
    unsafe = pure.is_absolute() or ".." in pure.parts or ":" in target or "\\" in target
    if unsafe:
        return 1.0, [f"链接目标越界或泄露绝对路径：{target!r}"], 1.0
    if target != XLSX_NAME:
        return 0.65, [f"链接指向非指定文件：{target!r}"], 1.0
    if relation.get("TargetMode") != "External" or not relation.get("Type", "").endswith("/hyperlink"):
        return 0.4, ["链接关系类型或 TargetMode 错误"], 1.0
    if not (artifact / target).is_file():
        return 0.75, ["链接目标文件未一同交付"], 1.0
    return 0.0, [], 1.0


def _extra_penalty(artifact: Path) -> tuple[float, list[str], float]:
    actual = {path.name for path in artifact.iterdir() if path.is_file()}
    directories = [path.name for path in artifact.iterdir() if path.is_dir()]
    extras = sorted(actual - set(REQUIRED)) + directories
    payloads = _zip_payloads(artifact / PPT_NAME)
    macro = any(name.endswith("vbaProject.bin") or name.endswith("vbaData.xml") for name in payloads)
    comments = any("comment" in name.casefold() for name in payloads)
    labels: list[str] = []
    if extras:
        labels.append(f"额外交付或子目录：{extras}")
    if macro:
        labels.append("PPT 含宏")
    if comments:
        labels.append("PPT 含批注")
    severity = min(1.0, 0.5 * bool(extras) + 0.75 * bool(macro) + 0.5 * bool(comments))
    return severity, labels, 1.0


def _safe_positive(call: Callable[[Path], tuple[float, list[str], int, int]], artifact: Path) -> tuple[float, list[str], float, float]:
    try:
        return call(artifact)
    except Exception as exc:
        return 0.0, [f"任务专属检查无法完成：{type(exc).__name__}: {exc}"], 0, 1


def _safe_penalty(call: Callable[[Path], tuple[float, list[str], float]], artifact: Path) -> tuple[float, list[str], float]:
    try:
        return call(artifact)
    except Exception as exc:
        return 1.0, [f"严重边界无法核验：{type(exc).__name__}: {exc}"], 1.0


def evaluate(artifact: Path) -> dict[str, Any]:
    ok, errors = _core(artifact)
    hashes = [_sha256(artifact / name) for name in REQUIRED if (artifact / name).is_file()]
    if not ok:
        return _fatal(errors, hashes)
    custom = {
        "R002": _workbook_positive,
        "R003": _button_positive,
        "R004": _geometry_positive,
        "R005": _relationship_positive,
        "R006": _preservation_positive,
        "R007": _render_positive,
        "R008": _style_positive,
        "R009": _security_positive,
    }
    results = [_criterion("R001", 1, "指定 PPTX/XLSX 存在、非空且可解析，未原样交回输入 PPT", 4, 4)]
    scores: dict[str, float] = {"R001": 1.0}
    for key in [f"R{index:03d}" for index in range(2, 10)]:
        score, failures, earned, total = _safe_positive(custom[key], artifact)
        results.append(_criterion(key, score, "本题任务专属检查通过" if not failures else "未通过：" + "；".join(failures), earned, total))
        scores[key] = score
    for key, call in (("P001", _unsafe_target_penalty), ("P002", _extra_penalty)):
        score, labels, total = _safe_penalty(call, artifact)
        results.append(_criterion(key, score, "未触发任务专属严重边界" if not labels else "；".join(labels), score * total, total))
        scores[key] = score
    positive_total = sum(POSITIVE_WEIGHTS.values())
    positive_earned = sum(POSITIVE_WEIGHTS[key] * scores[key] for key in POSITIVE_WEIGHTS)
    penalty_points = sum(abs(PENALTY_WEIGHTS[key]) * scores[key] for key in PENALTY_WEIGHTS)
    reward = round(max(0.0, min(1.0, (positive_earned - penalty_points) / positive_total)), 6)
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes,
        "criterion_results": results,
        "reward": reward,
        "passed": reward >= 0.8,
        "hurdle_failures": [],
        "penalties": [{"id": key, "amount": round(abs(PENALTY_WEIGHTS[key]) * scores[key] / positive_total, 6)} for key in PENALTY_WEIGHTS if scores[key] > 0],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    parser.add_argument("--scoring-version", choices=("v1", "v2"), default="v2")
    args = parser.parse_args()
    try:
        result = V1.evaluate(args.artifact_dir) if args.scoring_version == "v1" else evaluate(args.artifact_dir)
    except Exception as exc:
        result = _fatal([f"Judge 异常：{type(exc).__name__}: {exc}"], [])
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{float(result['reward']):.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
