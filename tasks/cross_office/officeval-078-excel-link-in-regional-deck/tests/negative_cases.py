#!/usr/bin/env python3
"""OfficeVal 078 的任务专属 v2 分层负例。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import shutil
import sys
import tempfile
import warnings
import zipfile
from pathlib import Path
from typing import Callable

from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution" / "Golden"
INPUTS = ROOT / "environment" / "input_files"
PPT_NAME = "区域协同汇报_第8页Excel超链接版.pptx"
SOURCE_PPT_NAME = "区域协同汇报.pptx"
XLSX_NAME = "CRM_pipeline_restyled_reordered.xlsx"
OUTPUTS = (PPT_NAME, XLSX_NAME)
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"p": P_NS, "a": A_NS, "r": R_NS, "rel": REL_NS}
Mutator = Callable[[Path], None]


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"无法载入 {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _rewrite_zip(path: Path, replacements: dict[str, bytes]) -> None:
    with zipfile.ZipFile(path) as source:
        entries = [(item, source.read(item.filename)) for item in source.infolist()]
    temporary = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(temporary, "w") as target:
        for item, payload in entries:
            target.writestr(item, replacements.get(item.filename, payload))
    temporary.replace(path)


def _roots(root: Path) -> tuple[etree._Element, etree._Element]:
    with zipfile.ZipFile(root / PPT_NAME) as archive:
        slide = etree.fromstring(archive.read("ppt/slides/slide8.xml"))
        relationships = etree.fromstring(archive.read("ppt/slides/_rels/slide8.xml.rels"))
    return slide, relationships


def _context(root: Path) -> tuple[etree._Element, etree._Element, etree._Element, etree._Element, etree._Element]:
    slide, relationships = _roots(root)
    link = slide.xpath(".//p:cNvPr/a:hlinkClick", namespaces=NS)[0]
    shape = link.xpath("ancestor::p:sp[1]", namespaces=NS)[0]
    relation_id = link.get(f"{{{R_NS}}}id")
    relation = relationships.xpath("./rel:Relationship[@Id=$rid]", namespaces=NS, rid=relation_id)[0]
    return slide, relationships, shape, link, relation


def _save(root: Path, slide: etree._Element, relationships: etree._Element) -> None:
    _rewrite_zip(root / PPT_NAME, {
        "ppt/slides/slide8.xml": etree.tostring(slide, xml_declaration=True, encoding="UTF-8", standalone=True),
        "ppt/slides/_rels/slide8.xml.rels": etree.tostring(relationships, xml_declaration=True, encoding="UTF-8", standalone=True),
    })


def _button_text(root: Path, first: str, second: str) -> None:
    slide, relationships, shape, _, _ = _context(root)
    nodes = shape.xpath(".//a:t", namespaces=NS)
    nodes[0].text = first
    nodes[1].text = second
    _save(root, slide, relationships)


def _accessible_name(root: Path) -> None:
    slide, relationships, shape, _, _ = _context(root)
    shape.xpath("./p:nvSpPr/p:cNvPr", namespaces=NS)[0].set("name", "普通按钮")
    _save(root, slide, relationships)


def _geometry(root: Path, element: str, attribute: str, centimeters: float) -> None:
    slide, relationships, shape, _, _ = _context(root)
    node = shape.xpath(f"./p:spPr/a:xfrm/a:{element}", namespaces=NS)[0]
    node.set(attribute, str(int(centimeters * 360000)))
    _save(root, slide, relationships)


def _relationship(root: Path, attribute: str, value: str | None) -> None:
    slide, relationships, _, _, relation = _context(root)
    if value is None:
        relation.attrib.pop(attribute, None)
    else:
        relation.set(attribute, value)
    _save(root, slide, relationships)


def _remove_relationship(root: Path) -> None:
    slide, relationships, _, _, relation = _context(root)
    relationships.remove(relation)
    _save(root, slide, relationships)


def _link_on_text_only(root: Path) -> None:
    slide, relationships, shape, link, _ = _context(root)
    relation_id = link.get(f"{{{R_NS}}}id")
    link.getparent().remove(link)
    run_properties = shape.xpath(".//a:r/a:rPr", namespaces=NS)[0]
    moved = etree.SubElement(run_properties, f"{{{A_NS}}}hlinkClick")
    moved.set(f"{{{R_NS}}}id", relation_id)
    _save(root, slide, relationships)


def _style(root: Path, xpath: str, attribute: str, value: str) -> None:
    slide, relationships, shape, _, _ = _context(root)
    node = shape.xpath(xpath, namespaces=NS)[0]
    node.set(attribute, value)
    _save(root, slide, relationships)


def _change_original_table(root: Path) -> None:
    slide, relationships, shape, _, _ = _context(root)
    nodes = [node for node in slide.xpath(".//a:t", namespaces=NS) if node.text == "HarborLink"]
    nodes[0].text = "HarborLink ERROR"
    _save(root, slide, relationships)


def _change_other_slide(root: Path) -> None:
    path = root / PPT_NAME
    deck = Presentation(path)
    for shape in deck.slides[1].shapes:
        if getattr(shape, "has_text_frame", False) and "Financial Snapshot" in shape.text:
            shape.text = shape.text.replace("Financial Snapshot", "Financial Snapshot ERROR", 1)
            deck.save(path)
            return
    raise ValueError("第2页未找到标题")


def _remove_last_slide(root: Path) -> None:
    path = root / PPT_NAME
    deck = Presentation(path)
    slide_id = deck.slides._sldIdLst[-1]
    deck.part.drop_rel(slide_id.rId)
    deck.slides._sldIdLst.remove(slide_id)
    deck.save(path)


def _workbook_cell(root: Path) -> None:
    path = root / XLSX_NAME
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        book = load_workbook(path)
    book["Strategic Pipeline"]["A4"] = "Wrong Owner"
    book.save(path)
    book.close()


def _workbook_sheet_removed(root: Path) -> None:
    path = root / XLSX_NAME
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        book = load_workbook(path)
    book.remove(book["Lead Development"])
    book.save(path)
    book.close()


def _input_only(root: Path) -> None:
    for path in list(root.iterdir()):
        if path.is_file():
            path.unlink()
    shutil.copy2(INPUTS / SOURCE_PPT_NAME, root / PPT_NAME)
    shutil.copy2(INPUTS / XLSX_NAME, root / XLSX_NAME)


def _severe_incomplete(root: Path) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Strategic Pipeline Status Update"
    deck.save(root / PPT_NAME)


def _cases() -> list[tuple[str, Mutator, str]]:
    cases: list[tuple[str, Mutator, str]] = []
    for name in OUTPUTS:
        cases.extend([
            (f"missing-{name}", lambda root, n=name: (root / n).unlink(), "fatal_zero"),
            (f"empty-{name}", lambda root, n=name: (root / n).write_bytes(b""), "fatal_zero"),
            (f"corrupt-{name}", lambda root, n=name: (root / n).write_bytes(b"not-a-valid-office-file"), "fatal_zero"),
        ])
    cases.extend([
        ("empty-output", lambda root: [path.unlink() for path in list(root.iterdir()) if path.is_file()], "fatal_zero"),
        ("unmodified-input-renamed", _input_only, "fatal_zero"),
        ("button-action-text-wrong", lambda root: _button_text(root, "查看文件", "CRM 工作簿"), "partial"),
        ("button-workbook-text-wrong", lambda root: _button_text(root, "打开 Excel", "普通文件"), "partial"),
        ("button-accessible-name-lost", _accessible_name, "partial"),
        ("button-moved-into-table", lambda root: _geometry(root, "off", "x", 23.0), "partial"),
        ("button-moved-into-title", lambda root: _geometry(root, "off", "y", 2.0), "partial"),
        ("button-too-narrow", lambda root: _geometry(root, "ext", "cx", 2.0), "partial"),
        ("button-too-tall", lambda root: _geometry(root, "ext", "cy", 6.0), "partial"),
        ("target-wrong-workbook", lambda root: _relationship(root, "Target", "Other.xlsx"), "partial"),
        ("target-absolute-path", lambda root: _relationship(root, "Target", "/tmp/CRM_pipeline_restyled_reordered.xlsx"), "partial"),
        ("target-parent-path", lambda root: _relationship(root, "Target", "../CRM_pipeline_restyled_reordered.xlsx"), "partial"),
        ("target-mode-missing", lambda root: _relationship(root, "TargetMode", None), "partial"),
        ("relationship-type-wrong", lambda root: _relationship(root, "Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"), "partial"),
        ("run-link-only", _link_on_text_only, "partial"),
        ("button-square-not-rounded", lambda root: _style(root, "./p:spPr/a:prstGeom", "prst", "rect"), "partial"),
        ("button-fill-red", lambda root: _style(root, "./p:spPr/a:solidFill/a:srgbClr", "val", "C62828"), "partial"),
        ("button-line-wrong", lambda root: _style(root, "./p:spPr/a:ln/a:solidFill/a:srgbClr", "val", "000000"), "partial"),
        ("button-text-dark", lambda root: _style(root, ".//a:rPr/a:solidFill/a:srgbClr", "val", "091426"), "partial"),
        ("original-pipeline-row-changed", _change_original_table, "partial"),
        ("other-slide-title-changed", _change_other_slide, "partial"),
        ("last-slide-removed", _remove_last_slide, "partial"),
        ("workbook-record-changed", _workbook_cell, "partial"),
        ("workbook-sheet-removed", _workbook_sheet_removed, "partial"),
        ("extra-debug-file", lambda root: (root / "debug.txt").write_text("debug", encoding="utf-8"), "partial"),
        ("moderate-visible-button-with-broken-link", _remove_relationship, "moderate"),
        ("severe-valid-one-slide-deck", _severe_incomplete, "severe"),
    ])
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    judge = _load("officeval_078_v2_judge", ROOT / "tests" / "judge.py")
    if judge.evaluate(GOLDEN).get("reward") != 1.0:
        raise AssertionError("Golden v2 必须为 1.0")
    results: list[dict] = []
    for name, mutator, layer in _cases():
        with tempfile.TemporaryDirectory(prefix=f"officeval-078-v2-{name}-") as temporary:
            output = Path(temporary) / "output"
            shutil.copytree(GOLDEN, output)
            mutator(output)
            result = judge.evaluate(output)
            reward = float(result.get("reward", 0))
            if layer == "fatal_zero" and reward != 0:
                raise AssertionError(f"{name} 应硬归零，实际 {reward}")
            if layer in {"partial", "moderate", "severe"} and not (0 < reward < 1):
                raise AssertionError(f"{name} 应保留连续部分分，实际 {reward}")
            if layer == "moderate" and reward >= 0.8:
                raise AssertionError(f"{name} 应低于通过线，实际 {reward}")
            results.append({"case": name, "layer": layer, "reward": reward, "passed": bool(result.get("passed")), "penalties": result.get("penalties", [])})
    moderate = next(item["reward"] for item in results if item["layer"] == "moderate")
    severe = next(item["reward"] for item in results if item["layer"] == "severe")
    if not severe < moderate:
        raise AssertionError(f"严重负例 {severe} 应低于中度负例 {moderate}")
    partial = [item["reward"] for item in results if item["layer"] != "fatal_zero"]
    report = {
        "total": len(results),
        "fatal_zero": sum(item["layer"] == "fatal_zero" for item in results),
        "partial_continuous": len(partial),
        "partial_reward_range": [min(partial), max(partial)],
        "moderate_reward": moderate,
        "severe_reward": severe,
        "cases": results,
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
