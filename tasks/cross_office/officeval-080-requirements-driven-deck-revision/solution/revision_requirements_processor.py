#!/usr/bin/env python3
"""从真实《要求.docx》和源 PPTX 动态提取 officeval_080 修订计划。"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from io import BytesIO
from pathlib import Path

from docx import Document
from PIL import Image
from pptx import Presentation


EMU_PER_CM = 360000.0
CN_DIGITS = {
    "零": 0,
    "一": 1,
    "二": 2,
    "两": 2,
    "三": 3,
    "四": 4,
    "五": 5,
    "六": 6,
    "七": 7,
    "八": 8,
    "九": 9,
}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def chinese_number(value: str) -> int:
    value = value.strip()
    if value.isdigit():
        return int(value)
    if value == "十":
        return 10
    if "十" in value:
        left, right = value.split("十", 1)
        tens = CN_DIGITS.get(left, 1) if left else 1
        ones = CN_DIGITS.get(right, 0) if right else 0
        return tens * 10 + ones
    if value in CN_DIGITS:
        return CN_DIGITS[value]
    raise ValueError(f"无法解析中文数字：{value!r}")


def page_mentions(text: str) -> list[int]:
    pages: list[int] = []
    for value in re.findall(r"第([零一二两三四五六七八九十\d]+)页", text):
        page = chinese_number(value)
        if page not in pages:
            pages.append(page)
    return pages


def first_matching(paragraphs: list[str], pattern: str) -> int:
    compiled = re.compile(pattern)
    for index, text in enumerate(paragraphs):
        if compiled.search(text):
            return index
    raise ValueError(f"要求文档中未找到模式：{pattern}")


def drawing_assets(document: Document) -> list[dict]:
    assets: list[dict] = []
    for paragraph_index, paragraph in enumerate(document.paragraphs):
        for drawing_index, drawing in enumerate(paragraph._p.xpath(".//w:drawing")):
            relationship_ids = drawing.xpath(".//a:blip/@r:embed")
            extents = drawing.xpath(".//wp:extent")
            if len(relationship_ids) != 1:
                continue
            relationship_id = relationship_ids[0]
            part = document.part.related_parts[relationship_id]
            width_cm = height_cm = None
            if extents:
                width_cm = int(extents[0].get("cx")) / EMU_PER_CM
                height_cm = int(extents[0].get("cy")) / EMU_PER_CM
            with Image.open(BytesIO(part.blob)) as image:
                pixel_width, pixel_height = image.size
            assets.append(
                {
                    "document_order": len(assets) + 1,
                    "paragraph_index": paragraph_index,
                    "drawing_index": drawing_index,
                    "relationship_id": relationship_id,
                    "part_name": Path(str(part.partname)).name,
                    "sha256": hashlib.sha256(part.blob).hexdigest(),
                    "size_bytes": len(part.blob),
                    "pixel_width": pixel_width,
                    "pixel_height": pixel_height,
                    "display_width_cm": round(width_cm, 6) if width_cm is not None else None,
                    "display_height_cm": round(height_cm, 6) if height_cm is not None else None,
                    "display_aspect_ratio": (
                        round(width_cm / height_cm, 6)
                        if width_cm is not None and height_cm
                        else None
                    ),
                }
            )
    return assets


def paragraph_asset_indexes(assets: list[dict], paragraph_index: int) -> list[int]:
    return [
        asset["document_order"]
        for asset in assets
        if asset["paragraph_index"] == paragraph_index
    ]


def width_value(text: str) -> float:
    match = re.search(r"图宽(?:要)?\s*(\d+(?:\.\d+)?)\s*(?:厘米|cm)", text, re.I)
    if not match:
        raise ValueError(f"未找到图片宽度：{text!r}")
    return float(match.group(1))


def font_inventory(slide) -> list[dict]:
    result: list[dict] = []
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False) or not shape.text.strip():
            continue
        sizes: list[float] = []
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.font.size is not None:
                    sizes.append(round(run.font.size.pt, 3))
        result.append(
            {
                "shape_id": shape.shape_id,
                "text": shape.text,
                "font_sizes_pt": sizes,
            }
        )
    return result


def extract_plan(requirements_path: Path, source_pptx_path: Path) -> dict:
    document = Document(requirements_path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
    assets = drawing_assets(document)
    presentation = Presentation(source_pptx_path)

    cover_index = first_matching(paragraphs, r"第一页是别人做的.*第二页")
    slide3_index = first_matching(paragraphs, r"第三页.*图宽")
    slide4_index = first_matching(paragraphs, r"第四页.*图片放左侧")
    slide5_title_index = first_matching(paragraphs, r"第五页.*What is the story")
    slide5_action_index = first_matching(paragraphs, r"加上两幅图")
    slide5_asset_index = next(
        asset["paragraph_index"]
        for asset in assets
        if asset["paragraph_index"] > slide5_action_index
    )
    slide6_index = first_matching(paragraphs, r"第六页.*箭头")
    animation_range_index = first_matching(paragraphs, r"第八页\s*[-—至]\s*第九页")
    animation_action_index = first_matching(paragraphs, r"小标题.*对应内容.*不同动画")
    slide10_title_index = first_matching(paragraphs, r"原来第十页删除.*Match five stages")
    slide10_asset_index = next(
        asset["paragraph_index"]
        for asset in assets
        if asset["paragraph_index"] > slide10_title_index
    )
    slide10_size_index = first_matching(paragraphs, r"图片放到标题下方.*图宽")
    slide11_index = first_matching(paragraphs, r"第十一页.*相同的出现动画")
    removal_index = first_matching(paragraphs, r"右下角图片删除")
    transitions_index = first_matching(paragraphs, r"所有页面.*不同的切换动画")

    slide3_asset_numbers = paragraph_asset_indexes(assets, slide3_index)
    slide5_asset_numbers = paragraph_asset_indexes(assets, slide5_asset_index)
    slide10_asset_numbers = paragraph_asset_indexes(assets, slide10_asset_index)
    if len(slide3_asset_numbers) != 1 or len(slide5_asset_numbers) != 2 or len(slide10_asset_numbers) != 1:
        raise ValueError("嵌入图片与审阅意见的对应关系不完整")

    slide3_width = width_value(paragraphs[slide3_index])
    slide3_asset = assets[slide3_asset_numbers[0] - 1]
    slide3_height = slide3_width / slide3_asset["display_aspect_ratio"]
    slide10_width = width_value(paragraphs[slide10_size_index])
    slide10_asset = assets[slide10_asset_numbers[0] - 1]
    slide10_height = slide10_width / slide10_asset["display_aspect_ratio"]

    first_and_second_pages = page_mentions(paragraphs[cover_index])
    if 1 not in first_and_second_pages or 2 not in first_and_second_pages:
        raise ValueError("封面修改段落没有同时指向第一页和第二页")
    animation_pages = page_mentions(paragraphs[animation_range_index])
    if animation_pages != [8, 9]:
        raise ValueError(f"动画页范围异常：{animation_pages}")
    removal_pages = page_mentions(paragraphs[removal_index])
    if len(removal_pages) != 1:
        raise ValueError("右下角图片删除要求必须只指向一页")

    requirements = [
        {
            "id": "R01",
            "slides": [1],
            "action": "increase_cover_fonts",
            "source_paragraph_indexes": [cover_index],
            "source_text": paragraphs[cover_index],
            "font_delta_pt": [2.0, 3.0],
        },
        {
            "id": "R02",
            "slides": [2],
            "action": "match_background",
            "source_paragraph_indexes": [cover_index],
            "source_text": paragraphs[cover_index],
            "reference_slide": 3,
        },
        {
            "id": "R03",
            "slides": page_mentions(paragraphs[slide3_index]),
            "action": "insert_review_image_right",
            "source_paragraph_indexes": [slide3_index],
            "source_text": paragraphs[slide3_index],
            "asset_indexes": slide3_asset_numbers,
            "width_cm": slide3_width,
            "height_cm_from_source_display_ratio": round(slide3_height, 4),
        },
        {
            "id": "R04",
            "slides": page_mentions(paragraphs[slide4_index]),
            "action": "image_left_text_right",
            "source_paragraph_indexes": [slide4_index],
            "source_text": paragraphs[slide4_index],
        },
        {
            "id": "R05",
            "slides": page_mentions(paragraphs[slide5_title_index]),
            "action": "insert_two_review_images",
            "source_paragraph_indexes": [slide5_title_index, slide5_action_index, slide5_asset_index],
            "source_text": "\n".join(
                paragraphs[index]
                for index in [slide5_title_index, slide5_action_index]
                if paragraphs[index]
            ),
            "asset_indexes": slide5_asset_numbers,
            "display_aspect_ratios": [
                assets[index - 1]["display_aspect_ratio"]
                for index in slide5_asset_numbers
            ],
        },
        {
            "id": "R06",
            "slides": page_mentions(paragraphs[slide6_index]),
            "action": "replace_arrow_answers_with_two_lines",
            "source_paragraph_indexes": [slide6_index],
            "source_text": paragraphs[slide6_index],
            "lines_per_arrow": 2,
        },
        {
            "id": "R07",
            "slides": [animation_pages[0]],
            "action": "animate_heading_and_content_with_distinct_effects",
            "source_paragraph_indexes": [animation_range_index, animation_action_index],
            "source_text": "\n".join([paragraphs[animation_range_index], paragraphs[animation_action_index]]),
        },
        {
            "id": "R08",
            "slides": [animation_pages[1]],
            "action": "animate_heading_and_content_with_distinct_effects",
            "source_paragraph_indexes": [animation_range_index, animation_action_index],
            "source_text": "\n".join([paragraphs[animation_range_index], paragraphs[animation_action_index]]),
        },
        {
            "id": "R09",
            "slides": [10],
            "action": "replace_slide_with_review_image",
            "source_paragraph_indexes": [slide10_title_index, slide10_asset_index, slide10_size_index],
            "source_text": "\n".join([paragraphs[slide10_title_index], paragraphs[slide10_size_index]]),
            "title": "Match five stages of the story with frames",
            "asset_indexes": slide10_asset_numbers,
            "width_cm": slide10_width,
            "height_cm_from_source_display_ratio": round(slide10_height, 4),
        },
        {
            "id": "R10",
            "slides": page_mentions(paragraphs[slide11_index]),
            "action": "animate_all_non_title_text_same_effect",
            "source_paragraph_indexes": [slide11_index],
            "source_text": paragraphs[slide11_index],
        },
        {
            "id": "R11",
            "slides": removal_pages,
            "action": "remove_bottom_right_image",
            "source_paragraph_indexes": [removal_index],
            "source_text": paragraphs[removal_index],
        },
        {
            "id": "R12",
            "slides": list(range(1, len(presentation.slides) + 1)),
            "action": "unique_transition_each_slide",
            "source_paragraph_indexes": [transitions_index],
            "source_text": paragraphs[transitions_index],
        },
    ]

    return {
        "schema_version": "1.0",
        "requirements_source": {
            "filename": requirements_path.name,
            "sha256": sha256_file(requirements_path),
            "size_bytes": requirements_path.stat().st_size,
        },
        "presentation_source": {
            "filename": source_pptx_path.name,
            "sha256": sha256_file(source_pptx_path),
            "size_bytes": source_pptx_path.stat().st_size,
            "slide_count": len(presentation.slides),
            "slide_width_cm": round(presentation.slide_width / EMU_PER_CM, 6),
            "slide_height_cm": round(presentation.slide_height / EMU_PER_CM, 6),
            "slide_1_font_inventory": font_inventory(presentation.slides[0]),
        },
        "raw_paragraphs": [
            {"index": index, "text": text}
            for index, text in enumerate(paragraphs)
        ],
        "embedded_assets": assets,
        "requirements": requirements,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--requirements", type=Path, required=True)
    parser.add_argument("--source-pptx", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    arguments = parser.parse_args()
    plan = extract_plan(arguments.requirements, arguments.source_pptx)
    arguments.output.parent.mkdir(parents=True, exist_ok=True)
    arguments.output.write_text(
        json.dumps(plan, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
