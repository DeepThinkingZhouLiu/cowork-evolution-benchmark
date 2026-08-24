#!/usr/bin/env python3
"""逐题完成 OmegaUse-OfficeVal officeval_080 的参考修订。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import shutil
import sys
import tempfile
from pathlib import Path

sys.dont_write_bytecode = True

from docx import Document
from lxml import etree
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt


P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
NSMAP = {"p": P_NS, "a": A_NS}
OUTPUT_PPTX = "公开课课件_修改完成.pptx"
OUTPUT_XLSX = "修订响应.xlsx"
OUTPUT_PROCESSOR = "revision_requirements_processor.py"
OUTPUT_PLAN = "revision_plan.json"


def load_processor(path: Path):
    specification = importlib.util.spec_from_file_location("revision_processor", path)
    if specification is None or specification.loader is None:
        raise RuntimeError("无法载入要求解析器")
    module = importlib.util.module_from_spec(specification)
    specification.loader.exec_module(module)
    return module


def find_text_shape(slide, needle: str, *, exact: bool = False):
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        text = shape.text.strip()
        if (exact and text == needle) or (not exact and needle in text):
            return shape
    raise ValueError(f"第 {slide.slide_id} 张幻灯片未找到文本：{needle}")


def set_shape_font(shape, size_pt: float, *, color: RGBColor | None = None, bold: bool | None = None) -> None:
    for paragraph in shape.text_frame.paragraphs:
        paragraph.font.size = Pt(size_pt)
        if color is not None:
            paragraph.font.color.rgb = color
        if bold is not None:
            paragraph.font.bold = bold
        for run in paragraph.runs:
            run.font.size = Pt(size_pt)
            if color is not None:
                run.font.color.rgb = color
            if bold is not None:
                run.font.bold = bold


def remove_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def set_background(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def embedded_blobs(requirements_path: Path) -> dict[int, tuple[str, bytes]]:
    document = Document(requirements_path)
    result: dict[int, tuple[str, bytes]] = {}
    order = 0
    for paragraph in document.paragraphs:
        for drawing in paragraph._p.xpath(".//w:drawing"):
            relationship_ids = drawing.xpath(".//a:blip/@r:embed")
            if len(relationship_ids) != 1:
                continue
            order += 1
            part = document.part.related_parts[relationship_ids[0]]
            result[order] = (Path(str(part.partname)).name, part.blob)
    return result


def extract_assets(requirements_path: Path, directory: Path) -> dict[int, Path]:
    output: dict[int, Path] = {}
    for index, (name, blob) in embedded_blobs(requirements_path).items():
        path = directory / f"asset-{index}-{name}"
        path.write_bytes(blob)
        output[index] = path
    return output


def apply_cover_fonts(presentation: Presentation) -> None:
    slide = presentation.slides[0]
    targets = {
        "Comics & Animation": 28,
        "Story Writing Lab": 44,
        "Build a story from panels, bubbles, and action.": 18,
        "observe visual clues": 19,
    }
    for text, size in targets.items():
        set_shape_font(find_text_shape(slide, text), size)


def apply_slide2_background(presentation: Presentation) -> None:
    slide = presentation.slides[1]
    set_background(slide, RGBColor(0xFF, 0xF7, 0xE8))
    colors = {
        "Writing": RGBColor(0x18, 0x31, 0x3C),
        "From comic strip": RGBColor(0x68, 0x74, 0x7A),
        "Turn pictures": RGBColor(0xE8, 0x8B, 0x24),
        "notice character": RGBColor(0x18, 0x31, 0x3C),
        "Today we write": RGBColor(0x68, 0x74, 0x7A),
        "02": RGBColor(0x8B, 0x91, 0x91),
    }
    for needle, color in colors.items():
        shape = find_text_shape(slide, needle)
        for paragraph in shape.text_frame.paragraphs:
            paragraph.font.color.rgb = color
            for run in paragraph.runs:
                run.font.color.rgb = color


def apply_slide3_image(presentation: Presentation, plan: dict, assets: dict[int, Path]) -> None:
    requirement = next(item for item in plan["requirements"] if item["id"] == "R03")
    slide = presentation.slides[2]
    for needle in [
        "Identify frames",
        "Organize the plot",
        "Complete a comic script",
        "Revise through self-check",
        "Review and apply passive voice",
    ]:
        shape = find_text_shape(slide, needle)
        shape.width = Cm(14.8)
        shape.height = Cm(1.25)
        set_shape_font(shape, 15)
    width = float(requirement["width_cm"])
    height = round(float(requirement["height_cm_from_source_display_ratio"]), 2)
    picture = slide.shapes.add_picture(
        str(assets[requirement["asset_indexes"][0]]),
        Cm(20.55),
        Cm(5.0),
        width=Cm(width),
        height=Cm(height),
    )
    picture.name = "审阅意见提供的学生学习场景"
    picture._element.nvPicPr.cNvPr.set("descr", "要求.docx 内嵌学生学习场景图片")


def apply_slide4_layout(presentation: Presentation) -> None:
    slide = presentation.slides[3]
    picture = next(shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
    picture.left = Cm(1.9)
    picture.top = Cm(4.0)
    picture.width = Cm(14.1)
    picture.height = Cm(13.3)
    positions = {
        "Warm-up: meet the explorer": (17.8, 1.07, 14.3, 1.14),
        "Observe the character": (17.9, 2.39, 14.0, 0.64),
        "What kind of story": (17.8, 4.19, 14.2, 1.27),
        "Who is the hero": (18.0, 6.73, 14.0, 3.81),
        "stronger plot": (18.0, 12.70, 14.0, 1.22),
    }
    for needle, (left, top, width, height) in positions.items():
        shape = find_text_shape(slide, needle)
        shape.left, shape.top = Cm(left), Cm(top)
        shape.width, shape.height = Cm(width), Cm(height)
    rule = slide.shapes[2]
    rule.left = Cm(17.8)
    rule.width = Cm(6.5)


def apply_slide5_images(presentation: Presentation, plan: dict, assets: dict[int, Path]) -> None:
    requirement = next(item for item in plan["requirements"] if item["id"] == "R05")
    slide = presentation.slides[4]
    tagline = find_text_shape(slide, "Use title words")
    tagline.top = Cm(16.35)
    tagline.left = Cm(2.54)
    tagline.width = Cm(26.2)
    tagline.height = Cm(1.05)
    target_height = 3.7
    placements = [(10.8, 12.1), (17.35, 12.1)]
    for asset_index, aspect_ratio, (left, top) in zip(
        requirement["asset_indexes"],
        requirement["display_aspect_ratios"],
        placements,
    ):
        width = target_height * float(aspect_ratio)
        picture = slide.shapes.add_picture(
            str(assets[asset_index]),
            Cm(left),
            Cm(top),
            width=Cm(width),
            height=Cm(target_height),
        )
        picture.name = f"审阅意见提供的故事图片{asset_index - 1}"
        picture._element.nvPicPr.cNvPr.set("descr", "要求.docx 内嵌故事图片")


def apply_slide6_lines(presentation: Presentation) -> None:
    slide = presentation.slides[5]
    arrow_shapes = [
        shape
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and "→" in shape.text
    ]
    if len(arrow_shapes) != 4:
        raise ValueError(f"第6页箭头条目数量异常：{len(arrow_shapes)}")
    for shape in arrow_shapes:
        shape.text = "→\n__________\n__________"
        shape.height = Cm(1.65)
        set_shape_font(shape, 11.5, color=RGBColor(0x18, 0x31, 0x3C))
        for paragraph in shape.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.LEFT


def clear_slide(slide) -> None:
    for shape in list(slide.shapes):
        remove_shape(shape)


def add_textbox(slide, text: str, left: float, top: float, width: float, height: float, size: float, *, bold: bool = False, color: RGBColor = RGBColor(0x18, 0x31, 0x3C)):
    shape = slide.shapes.add_textbox(Cm(left), Cm(top), Cm(width), Cm(height))
    shape.text = text
    shape.name = f"文本-{text[:24]}"
    set_shape_font(shape, size, color=color, bold=bold)
    return shape


def apply_slide10_replacement(presentation: Presentation, plan: dict, assets: dict[int, Path]) -> None:
    requirement = next(item for item in plan["requirements"] if item["id"] == "R09")
    slide = presentation.slides[9]
    clear_slide(slide)
    set_background(slide, RGBColor(0xFF, 0xF7, 0xE8))
    title = add_textbox(
        slide,
        requirement["title"],
        1.65,
        1.07,
        30.0,
        1.35,
        25,
        bold=True,
    )
    title.name = "第10页标题"
    accent = slide.shapes.add_shape(1, Cm(1.65), Cm(3.25), Cm(6.48), Cm(0.06))
    accent.name = "标题下方强调线"
    accent.fill.solid()
    accent.fill.fore_color.rgb = RGBColor(0x0D, 0x86, 0x8C)
    accent.line.fill.background()
    width = float(requirement["width_cm"])
    height = round(float(requirement["height_cm_from_source_display_ratio"]), 2)
    picture = slide.shapes.add_picture(
        str(assets[requirement["asset_indexes"][0]]),
        Cm((33.8666667 - width) / 2),
        Cm(3.85),
        width=Cm(width),
        height=Cm(height),
    )
    picture.name = "故事五阶段曲线图"
    picture._element.nvPicPr.cNvPr.set("descr", "要求.docx 内嵌故事五阶段曲线图")
    add_textbox(slide, "10", 31.29, 17.91, 1.14, 0.38, 8, color=RGBColor(0x8B, 0x91, 0x91))


def animation_behavior(parent, timing_id: int, shape_id: str, effect: str) -> int:
    child_list = etree.SubElement(parent, f"{{{P_NS}}}childTnLst")
    set_node = etree.SubElement(child_list, f"{{{P_NS}}}set")
    behavior = etree.SubElement(set_node, f"{{{P_NS}}}cBhvr")
    set_ctn = etree.SubElement(behavior, f"{{{P_NS}}}cTn", id=str(timing_id), dur="1", fill="hold")
    timing_id += 1
    conditions = etree.SubElement(set_ctn, f"{{{P_NS}}}stCondLst")
    etree.SubElement(conditions, f"{{{P_NS}}}cond", delay="0")
    target = etree.SubElement(behavior, f"{{{P_NS}}}tgtEl")
    etree.SubElement(target, f"{{{P_NS}}}spTgt", spid=shape_id)
    names = etree.SubElement(behavior, f"{{{P_NS}}}attrNameLst")
    name = etree.SubElement(names, f"{{{P_NS}}}attrName")
    name.text = "style.visibility"
    destination = etree.SubElement(set_node, f"{{{P_NS}}}to")
    etree.SubElement(destination, f"{{{P_NS}}}strVal", val="visible")

    effect_node = etree.SubElement(
        child_list,
        f"{{{P_NS}}}animEffect",
        transition="in",
        filter=effect,
    )
    effect_behavior = etree.SubElement(effect_node, f"{{{P_NS}}}cBhvr")
    etree.SubElement(effect_behavior, f"{{{P_NS}}}cTn", id=str(timing_id), dur="650", fill="hold")
    timing_id += 1
    effect_target = etree.SubElement(effect_behavior, f"{{{P_NS}}}tgtEl")
    etree.SubElement(effect_target, f"{{{P_NS}}}spTgt", spid=shape_id)
    return timing_id


def build_click_timing(targets: list[tuple[int, str, int]]) -> etree._Element:
    timing = etree.Element(f"{{{P_NS}}}timing", nsmap=NSMAP)
    time_nodes = etree.SubElement(timing, f"{{{P_NS}}}tnLst")
    root_parallel = etree.SubElement(time_nodes, f"{{{P_NS}}}par")
    root_ctn = etree.SubElement(
        root_parallel,
        f"{{{P_NS}}}cTn",
        id="1",
        dur="indefinite",
        restart="never",
        nodeType="tmRoot",
    )
    root_children = etree.SubElement(root_ctn, f"{{{P_NS}}}childTnLst")
    sequence = etree.SubElement(root_children, f"{{{P_NS}}}seq", concurrent="1", nextAc="seek")
    main_ctn = etree.SubElement(
        sequence,
        f"{{{P_NS}}}cTn",
        id="2",
        dur="indefinite",
        nodeType="mainSeq",
    )
    main_children = etree.SubElement(main_ctn, f"{{{P_NS}}}childTnLst")
    timing_id = 3
    for shape_id, effect, preset_id in targets:
        top_parallel = etree.SubElement(main_children, f"{{{P_NS}}}par")
        top_ctn = etree.SubElement(top_parallel, f"{{{P_NS}}}cTn", id=str(timing_id), fill="hold")
        timing_id += 1
        top_conditions = etree.SubElement(top_ctn, f"{{{P_NS}}}stCondLst")
        etree.SubElement(top_conditions, f"{{{P_NS}}}cond", delay="indefinite")
        top_children = etree.SubElement(top_ctn, f"{{{P_NS}}}childTnLst")
        behavior_parallel = etree.SubElement(top_children, f"{{{P_NS}}}par")
        behavior_ctn = etree.SubElement(behavior_parallel, f"{{{P_NS}}}cTn", id=str(timing_id), fill="hold")
        timing_id += 1
        behavior_conditions = etree.SubElement(behavior_ctn, f"{{{P_NS}}}stCondLst")
        etree.SubElement(behavior_conditions, f"{{{P_NS}}}cond", delay="0")
        behavior_children = etree.SubElement(behavior_ctn, f"{{{P_NS}}}childTnLst")
        effect_parallel = etree.SubElement(behavior_children, f"{{{P_NS}}}par")
        effect_ctn = etree.SubElement(
            effect_parallel,
            f"{{{P_NS}}}cTn",
            id=str(timing_id),
            presetID=str(preset_id),
            presetClass="entr",
            presetSubtype="0",
            fill="hold",
            nodeType="clickEffect",
        )
        timing_id += 1
        effect_conditions = etree.SubElement(effect_ctn, f"{{{P_NS}}}stCondLst")
        etree.SubElement(effect_conditions, f"{{{P_NS}}}cond", delay="0")
        timing_id = animation_behavior(effect_ctn, timing_id, str(shape_id), effect)

    previous = etree.SubElement(sequence, f"{{{P_NS}}}prevCondLst")
    previous_condition = etree.SubElement(previous, f"{{{P_NS}}}cond", evt="onPrev", delay="0")
    previous_target = etree.SubElement(previous_condition, f"{{{P_NS}}}tgtEl")
    etree.SubElement(previous_target, f"{{{P_NS}}}sldTgt")
    following = etree.SubElement(sequence, f"{{{P_NS}}}nextCondLst")
    following_condition = etree.SubElement(following, f"{{{P_NS}}}cond", evt="onNext", delay="0")
    following_target = etree.SubElement(following_condition, f"{{{P_NS}}}tgtEl")
    etree.SubElement(following_target, f"{{{P_NS}}}sldTgt")
    return timing


def set_timing(slide, timing: etree._Element) -> None:
    root = slide._element
    old = root.find(qn("p:timing"))
    if old is not None:
        root.remove(old)
    extension_list = root.find(qn("p:extLst"))
    if extension_list is None:
        root.append(timing)
    else:
        root.insert(root.index(extension_list), timing)


def apply_group_animations(slide, pairs: list[tuple[str, str]]) -> None:
    effects = [
        ("fade", 10),
        ("wipe(left)", 22),
        ("fly(inFromRight)", 2),
        ("checkerboard(across)", 5),
    ]
    targets: list[tuple[int, str, int]] = []
    for (heading, body), (effect, preset) in zip(pairs, effects):
        targets.append((find_text_shape(slide, heading, exact=True).shape_id, effect, preset))
        targets.append((find_text_shape(slide, body, exact=True).shape_id, effect, preset))
    set_timing(slide, build_click_timing(targets))


def apply_animations(presentation: Presentation) -> None:
    apply_group_animations(
        presentation.slides[7],
        [
            ("Speech bubble", "words spoken by characters"),
            ("Caption", "setting or scene information"),
            ("Sound effect", "noises in the scene"),
            ("Thought bubble", "words inside a character’s mind"),
        ],
    )
    apply_group_animations(
        presentation.slides[8],
        [
            ("Caption", "The scene begins at sunset."),
            ("Speech bubble", "“Watch out!”"),
            ("Thought bubble", "I hope the boat is strong."),
            ("Sound effect", "Crash!"),
        ],
    )
    slide11 = presentation.slides[10]
    subtitle = find_text_shape(slide11, "Task 1:")
    page_number = find_text_shape(slide11, "11", exact=True)
    dialogues = [
        shape
        for shape in slide11.shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip().startswith("“")
    ]
    dialogues.sort(key=lambda shape: (shape.top, shape.left))
    if len(dialogues) != 6:
        raise ValueError(f"第11页对白数量异常：{len(dialogues)}")
    targets = [(shape.shape_id, "fade", 10) for shape in [subtitle, page_number, *dialogues]]
    set_timing(slide11, build_click_timing(targets))


def apply_slide16_removal(presentation: Presentation, plan: dict) -> None:
    requirement = next(item for item in plan["requirements"] if item["id"] == "R11")
    page = int(requirement["slides"][0])
    slide = presentation.slides[page - 1]
    pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    candidates = [
        shape
        for shape in pictures
        if shape.left + shape.width / 2 > presentation.slide_width / 2
        and shape.top + shape.height / 2 > presentation.slide_height / 2
    ]
    if len(candidates) != 1:
        raise ValueError(f"第{page}页右下角图片数量异常：{len(candidates)}")
    remove_shape(candidates[0])


def add_transition(slide, effect_name: str, attributes: dict[str, str]) -> None:
    root = slide._element
    existing = root.find(qn("p:transition"))
    if existing is not None:
        root.remove(existing)
    transition = etree.Element(f"{{{P_NS}}}transition", spd="med", advClick="1")
    etree.SubElement(transition, f"{{{P_NS}}}{effect_name}", **attributes)
    timing = root.find(qn("p:timing"))
    extension = root.find(qn("p:extLst"))
    if timing is not None:
        root.insert(root.index(timing), transition)
    elif extension is not None:
        root.insert(root.index(extension), transition)
    else:
        root.append(transition)


def apply_transitions(presentation: Presentation) -> None:
    transitions = [
        ("fade", {}),
        ("push", {"dir": "l"}),
        ("push", {"dir": "r"}),
        ("push", {"dir": "u"}),
        ("push", {"dir": "d"}),
        ("wipe", {"dir": "l"}),
        ("wipe", {"dir": "r"}),
        ("wipe", {"dir": "u"}),
        ("wipe", {"dir": "d"}),
        ("cover", {"dir": "l"}),
        ("cover", {"dir": "r"}),
        ("cover", {"dir": "u"}),
        ("cover", {"dir": "d"}),
        ("split", {"orient": "horz", "dir": "in"}),
        ("split", {"orient": "horz", "dir": "out"}),
        ("split", {"orient": "vert", "dir": "in"}),
        ("split", {"orient": "vert", "dir": "out"}),
        ("wheel", {"spokes": "4"}),
    ]
    if len(presentation.slides) != len(transitions):
        raise ValueError("切换效果清单与页数不一致")
    for slide, (name, attributes) in zip(presentation.slides, transitions):
        add_transition(slide, name, attributes)


def build_response_workbook(plan: dict, output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "响应记录"
    headers = ["ID", "目标页", "原始审阅意见", "状态", "交叉文件核验"]
    sheet.append(headers)
    checks = {
        "R01": "pptx:slide=1;check=font_delta_2pt",
        "R02": "pptx:slide=2;check=background_match_slide3",
        "R03": "pptx:slide=3;check=embedded_image_size_position",
        "R04": "pptx:slide=4;check=image_left_text_right",
        "R05": "pptx:slide=5;check=two_embedded_images",
        "R06": "pptx:slide=6;check=two_lines_after_each_arrow",
        "R07": "pptx:slide=8;check=distinct_group_animations",
        "R08": "pptx:slide=9;check=distinct_group_animations",
        "R09": "pptx:slide=10;check=replaced_title_and_curve",
        "R10": "pptx:slide=11;check=same_click_animations",
        "R11": "pptx:slide=16;check=bottom_right_image_removed",
        "R12": "pptx:slides=1-18;check=unique_transitions",
    }
    for requirement in plan["requirements"]:
        slides = requirement["slides"]
        target = "1-18" if requirement["id"] == "R12" else ",".join(str(page) for page in slides)
        sheet.append(
            [
                requirement["id"],
                target,
                requirement["source_text"],
                "已完成",
                checks[requirement["id"]],
            ]
        )
    header_fill = PatternFill("solid", fgColor="0B3140")
    thin_border = Border(
        bottom=Side(style="thin", color="D8DFE3"),
    )
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = thin_border
    sheet.row_dimensions[1].height = 24
    for row_index in range(2, sheet.max_row + 1):
        source_text = str(sheet.cell(row_index, 3).value or "")
        sheet.row_dimensions[row_index].height = 48 if "\n" in source_text or len(source_text) > 58 else 30
    widths = [8, 10, 68, 12, 48]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:E{sheet.max_row}"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 2
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:1"
    sheet.print_area = f"A1:E{sheet.max_row}"
    workbook.save(output_path)


def solve(input_dir: Path, output_dir: Path) -> None:
    source_pptx = input_dir / "公开课课件.pptx"
    requirements = input_dir / "要求.docx"
    processor_path = Path(__file__).with_name("revision_requirements_processor.py")
    output_dir.mkdir(parents=True, exist_ok=True)
    processor = load_processor(processor_path)
    plan = processor.extract_plan(requirements, source_pptx)
    (output_dir / OUTPUT_PLAN).write_text(
        json.dumps(plan, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    shutil.copy2(processor_path, output_dir / OUTPUT_PROCESSOR)

    with tempfile.TemporaryDirectory(prefix="officeval-080-assets-") as temporary:
        assets = extract_assets(requirements, Path(temporary))
        presentation = Presentation(source_pptx)
        apply_cover_fonts(presentation)
        apply_slide2_background(presentation)
        apply_slide3_image(presentation, plan, assets)
        apply_slide4_layout(presentation)
        apply_slide5_images(presentation, plan, assets)
        apply_slide6_lines(presentation)
        apply_slide10_replacement(presentation, plan, assets)
        apply_animations(presentation)
        apply_slide16_removal(presentation, plan)
        apply_transitions(presentation)
        presentation.save(output_dir / OUTPUT_PPTX)

    build_response_workbook(plan, output_dir / OUTPUT_XLSX)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    arguments = parser.parse_args()
    solve(arguments.input_dir, arguments.output_dir)


if __name__ == "__main__":
    main()
