#!/usr/bin/env python3
"""OfficeVal 080 任务专属 v2 连续评分；原严格口径保存在 judge_v1.py。"""

from __future__ import annotations

import argparse
import ast
import importlib.util
import json
import math
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.dont_write_bytecode = True

TASK_ID = "cowork-evo/officeval-080-requirements-driven-deck-revision"
RUBRIC_VERSION = "omegause-officeval-080-adaptation-v0.2"
JUDGE_VERSION = "officeval-080-criterion-partial-v2.0"
OUTPUT_PPTX = "公开课课件_修改完成.pptx"
OUTPUT_XLSX = "修订响应.xlsx"
OUTPUT_PROCESSOR = "revision_requirements_processor.py"
OUTPUT_PLAN = "revision_plan.json"
REQUIRED_FILES = (OUTPUT_PPTX, OUTPUT_XLSX, OUTPUT_PROCESSOR, OUTPUT_PLAN)
POSITIVE_WEIGHTS = {
    "R001": 6,
    "R002": 12,
    "R003": 7,
    "R004": 8,
    "R005": 10,
    "R006": 6,
    "R007": 10,
    "R008": 8,
    "R009": 7,
    "R010": 8,
    "R011": 8,
    "R012": 10,
}
PENALTY_WEIGHTS = {"P001": -15, "P002": -4}


def _load_v1():
    path = Path(__file__).with_name("judge_v1.py")
    spec = importlib.util.spec_from_file_location("officeval_080_judge_v1", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入 v1 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


V1 = _load_v1()


def _criterion(key: str, score: float, evidence: str, earned: float, total: float) -> dict[str, Any]:
    return {
        "criterion_id": key,
        "score": round(max(0.0, min(1.0, score)), 6),
        "evidence": evidence,
        "method": "deterministic",
        "raw": {"earned_units": round(earned, 6), "total_units": round(total, 6)},
    }


def _ratio(findings: list[tuple[bool, str]]) -> tuple[float, str, int, int]:
    total = len(findings)
    earned = sum(bool(ok) for ok, _ in findings)
    failures = [label for ok, label in findings if not ok]
    evidence = "全部原子检查通过" if not failures else "未通过：" + "；".join(failures[:12])
    if len(failures) > 12:
        evidence += f"；另有 {len(failures) - 12} 项"
    return (earned / total if total else 0.0), evidence, earned, total


def _zip_ok(path: Path) -> bool:
    with zipfile.ZipFile(path) as archive:
        return archive.testzip() is None


def _core(artifact: Path, inputs: Path) -> tuple[bool, list[str]]:
    errors: list[str] = []
    if not artifact.is_dir():
        return False, ["交付目录不存在"]
    for name in REQUIRED_FILES:
        path = artifact / name
        if not path.is_file() or path.stat().st_size == 0:
            errors.append(f"核心交付缺失或为空：{name}")
    if errors:
        return False, errors
    try:
        deck = Presentation(artifact / OUTPUT_PPTX)
        if not deck.slides:
            raise ValueError("PPTX 无页面")
        if not _zip_ok(artifact / OUTPUT_PPTX):
            raise ValueError("PPTX ZIP 部件损坏")
    except Exception as error:
        errors.append(f"PPTX 无法解析：{error}")
    try:
        workbook = load_workbook(artifact / OUTPUT_XLSX, data_only=False)
        if not workbook.worksheets:
            raise ValueError("XLSX 无工作表")
        workbook.close()
        if not _zip_ok(artifact / OUTPUT_XLSX):
            raise ValueError("XLSX ZIP 部件损坏")
    except Exception as error:
        errors.append(f"XLSX 无法解析：{error}")
    try:
        plan = json.loads((artifact / OUTPUT_PLAN).read_text(encoding="utf-8"))
        if not isinstance(plan, dict):
            raise ValueError("计划根节点不是对象")
    except Exception as error:
        errors.append(f"revision_plan.json 无法解析：{error}")
    try:
        source = (artifact / OUTPUT_PROCESSOR).read_text(encoding="utf-8")
        ast.parse(source)
    except Exception as error:
        errors.append(f"处理器无法解析：{error}")
    try:
        if V1.sha256_file(artifact / OUTPUT_PPTX) == V1.sha256_file(inputs / "公开课课件.pptx"):
            errors.append("仅将未修改输入 PPTX 改名交回")
    except Exception as error:
        errors.append(f"无法核对源 PPTX：{error}")
    return not errors, errors


def _fatal(errors: list[str], artifact: Path) -> dict[str, Any]:
    hashes = []
    for name in REQUIRED_FILES:
        path = artifact / name
        if path.is_file() and path.stat().st_size:
            try:
                hashes.append(V1.sha256_file(path))
            except Exception:
                pass
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": hashes or ["0" * 64],
        "criterion_results": [
            _criterion(key, 0, "；".join(errors), 0, 0)
            for key in (*POSITIVE_WEIGHTS, *PENALTY_WEIGHTS)
        ],
        "score": 0.0,
        "reward": 0.0,
        "passed": False,
        "hurdle_failures": ["R001"],
        "penalties": [],
    }


def _plan_processor(artifact: Path, inputs: Path, truth: dict) -> tuple[float, str, int, int, dict[str, Any]]:
    plan = json.loads((artifact / OUTPUT_PLAN).read_text(encoding="utf-8"))
    try:
        mapped = V1.requirement_map(plan)
    except Exception:
        mapped = {}
    findings: list[tuple[bool, str]] = [
        (plan.get("requirements_source", {}).get("sha256") == truth["requirements_sha256"], "要求文档哈希错误"),
        (plan.get("presentation_source", {}).get("sha256") == truth["source_pptx_sha256"], "源 PPTX 哈希错误"),
        (plan.get("presentation_source", {}).get("slide_count") == truth["slide_count"], "源 PPTX 页数错误"),
        ([item.get("text") for item in plan.get("raw_paragraphs", [])] == truth["paragraphs"], "原始段落未忠实保存"),
        ([item.get("sha256") for item in plan.get("embedded_assets", [])] == [item["sha256"] for item in truth["assets"]], "四张内嵌图片顺序或哈希错误"),
        (set(mapped) == set(V1.EXPECTED_ACTIONS), "R01–R12 ID 不完整或重复"),
    ]
    for identifier, action in V1.EXPECTED_ACTIONS.items():
        item = mapped.get(identifier, {})
        findings.extend([
            (item.get("action") == action, f"{identifier} action 错误"),
            (item.get("slides") == truth["expected_slides"][identifier], f"{identifier} 目标页错误"),
            (item.get("source_text") == truth["expected_text"][identifier], f"{identifier} 原始意见错误"),
        ])
    for identifier, width_key, height_key in [
        ("R03", "slide3_width", "slide3_height"),
        ("R09", "slide10_width", "slide10_height"),
    ]:
        item = mapped.get(identifier, {})
        try:
            width_ok = math.isclose(float(item.get("width_cm", -1)), truth[width_key], abs_tol=1e-6)
            height_ok = math.isclose(float(item.get("height_cm_from_source_display_ratio", -1)), truth[height_key], abs_tol=0.01)
        except (TypeError, ValueError):
            width_ok = height_ok = False
        findings.extend([(width_ok, f"{identifier} 动态宽度错误"), (height_ok, f"{identifier} 动态高度错误")])

    replay = {"canonical": False, "metamorphic": False, "details": []}
    processor = artifact / OUTPUT_PROCESSOR
    with tempfile.TemporaryDirectory(prefix="officeval-080-v2-replay-") as temporary:
        root = Path(temporary)
        canonical = root / "canonical.json"
        mutated_doc = root / "要求_内容变异.docx"
        mutated_plan_path = root / "mutated.json"
        try:
            ok, detail = V1.run_processor(processor, inputs / "要求.docx", inputs / "公开课课件.pptx", canonical)
            replay["details"].append(detail)
            canonical_valid = False
            if ok:
                generated = json.loads(canonical.read_text(encoding="utf-8"))
                canonical_valid, errors = V1.validate_plan(generated, truth)
                replay["details"].extend(errors)
            replay["canonical"] = bool(ok and canonical_valid)
        except Exception as error:
            replay["details"].append(f"原输入 replay：{error}")
        findings.append((replay["canonical"], "原输入处理器 replay 未通过"))
        try:
            V1.mutate_requirements(inputs / "要求.docx", mutated_doc)
            mutated_truth = V1.input_truth(mutated_doc, inputs / "公开课课件.pptx")
            ok, detail = V1.run_processor(processor, mutated_doc, inputs / "公开课课件.pptx", mutated_plan_path)
            replay["details"].append(detail)
            changed = False
            mutated_valid = False
            if ok:
                generated = json.loads(mutated_plan_path.read_text(encoding="utf-8"))
                mutated_valid, errors = V1.validate_plan(generated, mutated_truth)
                replay["details"].extend(errors)
                mapped_mutated = V1.requirement_map(generated)
                changed = (
                    math.isclose(float(mapped_mutated.get("R03", {}).get("width_cm", -1)), 13.0, abs_tol=1e-6)
                    and mapped_mutated.get("R11", {}).get("slides") == [17]
                    and "图宽13厘米" in mapped_mutated.get("R03", {}).get("source_text", "")
                    and "第十七页" in mapped_mutated.get("R11", {}).get("source_text", "")
                )
            replay["metamorphic"] = bool(ok and mutated_valid and changed)
        except Exception as error:
            replay["details"].append(f"内容变异 replay：{error}")
        findings.append((replay["metamorphic"], "12cm→13cm、16页→17页内容变异 replay 未通过"))
    score, evidence, earned, total = _ratio(findings)
    return score, evidence, earned, total, replay


def _cover_background(source, output) -> tuple[float, str, int, int]:
    findings: list[tuple[bool, str]] = []
    for anchor in [
        "Comics & Animation",
        "Story Writing Lab",
        "Build a story from panels, bubbles, and action.",
        "observe visual clues",
    ]:
        before = V1.font_sizes(V1.get_shape(source.slides[0], anchor))
        after = V1.font_sizes(V1.get_shape(output.slides[0], anchor))
        ok = bool(before and after and len(before) == len(after)) and all(2 <= right - left <= 3 for left, right in zip(before, after))
        findings.append((ok, f"封面 {anchor!r} 未逐 run 增加 2–3pt"))
    bg2, bg3 = V1.background_rgb(output.slides[1]), V1.background_rgb(output.slides[2])
    findings.extend([(bg2 == bg3, f"第2页背景 {bg2} 未匹配第3页 {bg3}"), (bg2 == "FFF7E8", f"第2页背景不是米黄色：{bg2}")])
    return _ratio(findings)


def _slide3(output, truth) -> tuple[float, str, int, int]:
    slide = output.slides[2]
    expected_hash = truth["assets"][0]["sha256"]
    pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and V1.picture_hash(shape) == expected_hash]
    picture = pictures[0] if len(pictures) == 1 else None
    text_shapes = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()]
    overlap = max((V1.overlap_fraction(picture, shape) for shape in text_shapes), default=1.0) if picture is not None else 1.0
    findings = [
        (len(pictures) == 1, f"真实学生图数量应为1，实际{len(pictures)}"),
        (picture is not None and abs(picture.width / V1.EMU_PER_CM - truth["slide3_width"]) <= 0.05, "学生图宽度不是动态 12cm"),
        (picture is not None and abs(picture.height / V1.EMU_PER_CM - truth["slide3_height"]) <= 0.05, "学生图高度未按 Word 显示比例重算"),
        (picture is not None and V1.center_x(picture) > output.slide_width / 2, "学生图未置于右侧"),
        (picture is not None and picture.left >= output.slide_width * 0.55, "学生图侵入左侧教学内容区"),
        (overlap < 0.05, f"学生图与文字重叠比例过大：{overlap:.3f}"),
        (picture is not None and picture.left + picture.width <= output.slide_width and picture.top + picture.height <= output.slide_height, "学生图越出页面"),
    ]
    return _ratio(findings)


def _slides4_5(source, output, truth) -> tuple[float, str, int, int]:
    slide4 = output.slides[3]
    source_picture = next((shape for shape in source.slides[3].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE), None)
    expected4 = V1.picture_hash(source_picture) if source_picture is not None else ""
    pictures4 = [shape for shape in slide4.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and V1.picture_hash(shape) == expected4]
    picture4 = pictures4[0] if len(pictures4) == 1 else None
    right_text = [V1.get_shape(slide4, needle) for needle in ["Warm-up: meet the explorer", "What kind of story", "Who is the hero", "stronger plot"]]
    findings: list[tuple[bool, str]] = [
        (len(pictures4) == 1, "第4页未保留唯一真实探险图"),
        (picture4 is not None and V1.center_x(picture4) < output.slide_width / 2, "第4页图片不在左侧"),
    ]
    for needle, shape in zip(["Warm-up", "story type", "hero", "plot"], right_text):
        findings.append((shape is not None, f"第4页缺少 {needle} 文字组"))
        findings.append((shape is not None and V1.center_x(shape) > output.slide_width / 2, f"第4页 {needle} 未在右侧"))
        findings.append((picture4 is not None and shape is not None and picture4.left + picture4.width <= shape.left, f"第4页图片与 {needle} 横向冲突"))

    slide5 = output.slides[4]
    pictures5 = [shape for shape in slide5.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    text5 = [shape for shape in slide5.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()]
    for asset_index in [1, 2]:
        asset = truth["assets"][asset_index]
        matches = [shape for shape in pictures5 if V1.picture_hash(shape) == asset["sha256"]]
        picture = matches[0] if len(matches) == 1 else None
        overlap = max((V1.overlap_fraction(picture, text) for text in text5), default=1.0) if picture is not None else 1.0
        findings.extend([
            (len(matches) == 1, f"第5页 Word 图片{asset_index + 1} 缺失或重复"),
            (picture is not None and abs(picture.width / picture.height - asset["display_ratio"]) <= 0.02, f"第5页图片{asset_index + 1} 比例错误"),
            (overlap < 0.05, f"第5页图片{asset_index + 1} 遮挡文字"),
            (picture is not None and picture.left + picture.width <= output.slide_width and picture.top + picture.height <= output.slide_height, f"第5页图片{asset_index + 1} 越界"),
        ])
    findings.append((len(pictures5) == 2, f"第5页应恰有两图，实际{len(pictures5)}"))
    return _ratio(findings)


def _slide6(output) -> tuple[float, str, int, int]:
    slide = output.slides[5]
    arrows = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False) and "→" in shape.text]
    findings: list[tuple[bool, str]] = [(len(arrows) == 4, f"箭头条目应为4，实际{len(arrows)}")]
    for index in range(4):
        if index < len(arrows):
            right = arrows[index].text.split("→", 1)[1]
            ok = len(__import__("re").findall(r"_+", right)) == 2 and __import__("re").sub(r"_|\s", "", right) == ""
        else:
            ok = False
        findings.append((ok, f"第{index + 1}个箭头后不是恰好两条空横线"))
    all_text = "\n".join(shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False))
    for answer in ["start the trip", "face danger", "make a choice", "find a way out"]:
        findings.append((answer not in all_text, f"旧答案仍存在：{answer}"))
    return _ratio(findings)


ANIMATION_GROUPS = {
    8: [
        ("Speech bubble", "words spoken by characters"),
        ("Caption", "setting or scene information"),
        ("Sound effect", "noises in the scene"),
        ("Thought bubble", "words inside a character’s mind"),
    ],
    9: [
        ("Caption", "The scene begins at sunset."),
        ("Speech bubble", "“Watch out!”"),
        ("Thought bubble", "I hope the boat is strong."),
        ("Sound effect", "Crash!"),
    ],
}


def _slides8_9(output) -> tuple[float, str, int, int]:
    findings: list[tuple[bool, str]] = []
    for page, pairs in ANIMATION_GROUPS.items():
        slide = output.slides[page - 1]
        effects = V1.effect_targets(slide)
        clicks = V1.click_targets(slide)
        signatures: list[tuple[str, ...]] = []
        for heading, body in pairs:
            heading_shape = V1.get_shape(slide, heading, exact=True)
            body_shape = V1.get_shape(slide, body, exact=True)
            heading_effect = effects.get(str(heading_shape.shape_id), set()) if heading_shape is not None else set()
            body_effect = effects.get(str(body_shape.shape_id), set()) if body_shape is not None else set()
            findings.extend([
                (heading_shape is not None, f"P{page} 缺少标题 {heading}"),
                (body_shape is not None, f"P{page} 缺少正文 {body}"),
                (len(heading_effect) == 1 and heading_effect == body_effect, f"P{page} {heading} 标题/正文进入动画不一致"),
                (heading_shape is not None and str(heading_shape.shape_id) in clicks, f"P{page} {heading} 标题未独立单击触发"),
                (body_shape is not None and str(body_shape.shape_id) in clicks, f"P{page} {heading} 正文未独立单击触发"),
            ])
            signatures.append(tuple(sorted(heading_effect)))
        findings.append((len(set(signatures)) == 4 and all(signatures), f"P{page} 四组动画效果不互异：{signatures}"))
    return _ratio(findings)


def _slide10(output, truth) -> tuple[float, str, int, int]:
    slide = output.slides[9]
    title = V1.get_shape(slide, "Match five stages of the story with frames", exact=True)
    expected_hash = truth["assets"][3]["sha256"]
    pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and V1.picture_hash(shape) == expected_hash]
    curve = pictures[0] if len(pictures) == 1 else None
    findings = [
        (title is not None, "第10页指定标题缺失或不精确"),
        (len(pictures) == 1, f"真实曲线图数量应为1，实际{len(pictures)}"),
        (curve is not None and abs(curve.width / V1.EMU_PER_CM - truth["slide10_width"]) <= 0.01, "曲线图宽度不是20cm"),
        (curve is not None and abs(curve.height / V1.EMU_PER_CM - truth["slide10_height"]) <= 0.01, "曲线图高度未按 Word 比例计算"),
        (curve is not None and title is not None and curve.top >= title.top + title.height, "曲线图未位于标题下方"),
        (curve is not None and curve.left >= 0 and curve.left + curve.width <= output.slide_width and curve.top + curve.height <= output.slide_height, "曲线图越界"),
    ]
    return _ratio(findings)


def _slide11(output) -> tuple[float, str, int, int]:
    slide = output.slides[10]
    shapes = [V1.get_shape(slide, "Task 1:"), V1.get_shape(slide, "11", exact=True)]
    dialogues = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip().startswith("“")]
    dialogues.sort(key=lambda shape: (shape.top, shape.left))
    shapes.extend(dialogues)
    effects = V1.effect_targets(slide)
    clicks = V1.click_targets(slide)
    findings: list[tuple[bool, str]] = [
        (len(shapes) == 8 and all(shape is not None for shape in shapes), f"应定位8个动画对象，实际{len([s for s in shapes if s is not None])}"),
        (len(dialogues) == 6, f"对白应为6条，实际{len(dialogues)}"),
    ]
    expected_ids = [str(shape.shape_id) for shape in shapes if shape is not None]
    for index, shape in enumerate(shapes):
        identifier = str(shape.shape_id) if shape is not None else "missing"
        findings.append((shape is not None and effects.get(identifier, set()) == {"fade"}, f"第{index + 1}个对象未使用 fade 进入动画"))
        findings.append((shape is not None and index < len(clicks) and clicks[index] == identifier, f"第{index + 1}个对象未按阅读顺序单击触发"))
    findings.append((clicks == expected_ids, f"完整点击序列错误：{clicks}"))
    return _ratio(findings)


def _slide16_transitions(output) -> tuple[float, str, int, int]:
    findings: list[tuple[bool, str]] = []
    if len(output.slides) >= 16:
        slide16 = output.slides[15]
        lower_right = [shape for shape in slide16.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and V1.center_x(shape) > output.slide_width / 2 and shape.top + shape.height / 2 > output.slide_height / 2]
        findings.append((not lower_right, f"第16页右下角仍有{len(lower_right)}张图片"))
    else:
        findings.append((False, "缺少第16页"))
    signatures = [V1.transition_signature(slide) for slide in output.slides]
    for page in range(18):
        findings.append((page < len(signatures) and signatures[page] is not None, f"第{page + 1}页缺少切换"))
    nonempty = [item for item in signatures[:18] if item is not None]
    findings.append((len(signatures) == 18 and len(set(nonempty)) == 18, f"切换唯一性为{len(set(nonempty))}/18"))
    return _ratio(findings)


def _response(artifact: Path, truth: dict, content_scores: dict[str, float]) -> tuple[float, str, int, int, bool]:
    workbook = load_workbook(artifact / OUTPUT_XLSX, data_only=False)
    findings: list[tuple[bool, str]] = [
        (workbook.sheetnames == ["响应记录"], f"工作表名称/数量错误：{workbook.sheetnames}"),
    ]
    sheet = workbook["响应记录"] if "响应记录" in workbook.sheetnames else workbook.worksheets[0]
    findings.extend([
        (sheet.max_row == 13, f"响应表行数应为13，实际{sheet.max_row}"),
        (sheet.max_column == 5, f"响应表列数应为5，实际{sheet.max_column}"),
        ([sheet.cell(1, column).value for column in range(1, 6)] == ["ID", "目标页", "原始审阅意见", "状态", "交叉文件核验"], "响应表表头错误"),
    ])
    rows: dict[str, list[Any]] = {}
    unsupported = False
    for row in range(2, sheet.max_row + 1):
        values = [sheet.cell(row, column).value for column in range(1, 6)]
        identifier = values[0]
        if isinstance(identifier, str) and identifier in V1.EXPECTED_ACTIONS and identifier not in rows:
            rows[identifier] = values
        elif any(value not in (None, "") for value in values):
            unsupported = True
    for identifier in V1.EXPECTED_ACTIONS:
        values = rows.get(identifier)
        findings.append((values is not None, f"响应表遗漏 {identifier}"))
        if values is None:
            continue
        _, target, source_text, status, check = values
        expected_target = "1-18" if identifier == "R12" else ",".join(str(page) for page in truth["expected_slides"][identifier])
        findings.extend([
            (str(target) == expected_target, f"{identifier} 目标页错误"),
            (source_text == truth["expected_text"][identifier], f"{identifier} 原始审阅意见错误"),
            (status == "已完成", f"{identifier} 状态不是已完成"),
            (check == V1.EXPECTED_RESPONSE_CHECKS[identifier], f"{identifier} 交叉核验字段错误"),
        ])
        criterion_map = {"R01": "R003", "R02": "R003", "R03": "R004", "R04": "R005", "R05": "R005", "R06": "R006", "R07": "R007", "R08": "R007", "R09": "R008", "R10": "R009", "R11": "R010", "R12": "R010"}
        referenced_score = content_scores.get(criterion_map[identifier], 0)
        findings.append((status != "已完成" or referenced_score >= 0.5, f"{identifier} 声称完成但对应 PPT 完成度低于50%"))
    workbook.close()
    score, evidence, earned, total = _ratio(findings)
    return score, evidence, earned, total, unsupported


def _professional(source, output, artifact: Path, run_render: bool) -> tuple[float, str, int, int]:
    findings: list[tuple[bool, str]] = [
        (len(output.slides) == 18, f"PPT 页数应为18，实际{len(output.slides)}"),
    ]
    total_text = sum(1 for slide in output.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip())
    findings.append((total_text >= 120, f"原生文本对象过少：{total_text}"))
    for page in list(range(1, 10)) + list(range(11, 19)):
        if page > len(source.slides) or page > len(output.slides):
            findings.append((False, f"P{page} 不存在，无法核对标题"))
            continue
        title = next((shape.text.strip() for shape in source.slides[page - 1].shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()), None)
        findings.append((title is None or V1.get_shape(output.slides[page - 1], title, exact=True) is not None, f"P{page} 原标题未保留"))
    out_of_bounds = 0
    overlaps = 0
    for slide in output.slides:
        text_shapes = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()]
        for shape in text_shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > output.slide_width or shape.top + shape.height > output.slide_height:
                out_of_bounds += 1
        for index, first in enumerate(text_shapes):
            for second in text_shapes[index + 1:]:
                if V1.overlap_fraction(first, second) > 0.20:
                    overlaps += 1
    findings.extend([(out_of_bounds == 0, f"文字越界{out_of_bounds}处"), (overlaps < 3, f"显著文字重叠{overlaps}处")])
    try:
        upstream_path = Path(__file__).resolve().parents[1] / "adaptation" / "upstream-verifier.py"
        spec = importlib.util.spec_from_file_location("officeval_080_upstream_v2", upstream_path)
        module = importlib.util.module_from_spec(spec)
        assert spec is not None and spec.loader is not None
        spec.loader.exec_module(module)
        upstream = module.evaluate(str(artifact))
        findings.append((upstream.get("total_score") == 31 and upstream.get("max_score") == 31, f"冻结上游 verifier={upstream.get('total_score')}/{upstream.get('max_score')}"))
    except Exception as error:
        findings.append((False, f"冻结上游 verifier 异常：{error}"))
    if run_render:
        try:
            ok, detail = V1.libreoffice_render_check(artifact / OUTPUT_PPTX, artifact / OUTPUT_XLSX)
        except Exception as error:
            ok, detail = False, str(error)
        findings.append((ok, f"LibreOffice 渲染失败：{detail}"))
    else:
        findings.append((True, "分层负例快速测试跳过渲染"))
    with zipfile.ZipFile(artifact / OUTPUT_PPTX) as archive:
        macro_ppt = any(name.lower().endswith("vbaproject.bin") for name in archive.namelist())
    with zipfile.ZipFile(artifact / OUTPUT_XLSX) as archive:
        macro_xlsx = any(name.lower().endswith("vbaproject.bin") for name in archive.namelist())
    findings.append((not macro_ppt and not macro_xlsx, "交付包含宏"))
    return _ratio(findings)


def evaluate(artifact_dir: Path, input_dir: Path | None = None, *, run_render: bool = True) -> dict[str, Any]:
    artifact = Path(artifact_dir)
    inputs = Path(input_dir) if input_dir is not None else Path(__file__).resolve().parents[1] / "environment" / "input_files"
    core_ok, core_errors = _core(artifact, inputs)
    if not core_ok:
        return _fatal(core_errors, artifact)
    truth = V1.input_truth(inputs / "要求.docx", inputs / "公开课课件.pptx")
    source = Presentation(inputs / "公开课课件.pptx")
    output = Presentation(artifact / OUTPUT_PPTX)
    results: dict[str, tuple[float, str, int, int]] = {}
    results["R001"] = (1.0, "四个核心交付可解析且 PPTX 不是原样输入", 4, 4)
    plan_score, plan_evidence, plan_earned, plan_total, replay = _plan_processor(artifact, inputs, truth)
    results["R002"] = (plan_score, plan_evidence, plan_earned, plan_total)
    def safe(key: str, callback) -> None:
        try:
            results[key] = callback()
        except Exception as error:
            results[key] = (0.0, f"该项无法检查：{error}", 0, 1)

    safe("R003", lambda: _cover_background(source, output))
    safe("R004", lambda: _slide3(output, truth))
    safe("R005", lambda: _slides4_5(source, output, truth))
    safe("R006", lambda: _slide6(output))
    safe("R007", lambda: _slides8_9(output))
    safe("R008", lambda: _slide10(output, truth))
    safe("R009", lambda: _slide11(output))
    safe("R010", lambda: _slide16_transitions(output))
    content_scores = {key: value[0] for key, value in results.items()}
    response_score, response_evidence, response_earned, response_total, unsupported = _response(artifact, truth, content_scores)
    results["R011"] = (response_score, response_evidence, response_earned, response_total)
    safe("R012", lambda: _professional(source, output, artifact, run_render))

    processor_source = (artifact / OUTPUT_PROCESSOR).read_text(encoding="utf-8")
    hardcoded = ("with_name('revision_plan.json')" in processor_source or 'with_name("revision_plan.json")' in processor_source) and not replay["metamorphic"]
    false_completion = response_score > 0.85 and sum(score < 0.5 for key, (score, _, _, _) in results.items() if key in {"R003", "R004", "R005", "R006", "R007", "R008", "R009", "R010"}) >= 4
    p1_triggered = unsupported or hardcoded or false_completion
    p1_reasons = [label for ok, label in [(unsupported, "响应表含无来源行"), (hardcoded, "处理器直接复制固定计划且变异 replay 失败"), (false_completion, "大量修订未完成但响应表仍整表声明完成")] if ok]
    results["P001"] = (1.0 if p1_triggered else 0.0, "；".join(p1_reasons) if p1_reasons else "未触发", int(p1_triggered), 1)
    allowed = set(REQUIRED_FILES) | {"solve.py", "solve.sh"}
    extras = sorted(path.name for path in artifact.iterdir() if path.is_file() and path.name not in allowed)
    results["P002"] = (1.0 if extras else 0.0, f"额外文件={extras}" if extras else "未触发", int(bool(extras)), 1)

    positive = sum(POSITIVE_WEIGHTS[key] * results[key][0] for key in POSITIVE_WEIGHTS) / sum(POSITIVE_WEIGHTS.values())
    penalties = sum((-PENALTY_WEIGHTS[key]) * results[key][0] for key in PENALTY_WEIGHTS) / sum(POSITIVE_WEIGHTS.values())
    reward = round(max(0.0, min(1.0, positive - penalties)), 6)
    criterion_results = [
        _criterion(key, results[key][0], results[key][1], results[key][2], results[key][3])
        for key in (*POSITIVE_WEIGHTS, *PENALTY_WEIGHTS)
    ]
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [V1.sha256_file(artifact / name) for name in REQUIRED_FILES],
        "criterion_results": criterion_results,
        "score": reward,
        "reward": reward,
        "passed": reward >= 0.8,
        "hurdle_failures": [],
        "penalties": [key for key in PENALTY_WEIGHTS if results[key][0] > 0],
        "metamorphic_replay": replay,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--input-dir", type=Path, default=Path("/app/data/input_files"))
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--skip-render", action="store_true")
    parser.add_argument("--scoring-version", choices=["v1", "v2"], default="v2")
    arguments = parser.parse_args()
    if arguments.scoring_version == "v1":
        result = V1.evaluate(arguments.artifact_dir, arguments.input_dir, run_render=not arguments.skip_render)
        result["reward"] = result.get("score", 0.0)
        result["scoring_version"] = "v1"
    else:
        result = evaluate(arguments.artifact_dir, arguments.input_dir, run_render=not arguments.skip_render)
        result["scoring_version"] = "v2"
    arguments.result.parent.mkdir(parents=True, exist_ok=True)
    arguments.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
