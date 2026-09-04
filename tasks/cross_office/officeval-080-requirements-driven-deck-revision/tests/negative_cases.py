#!/usr/bin/env python3
"""OfficeVal 080 的任务专属 v2 分层负例。"""

from __future__ import annotations

import argparse
import importlib.util
import json
import shutil
import sys
import tempfile
import zipfile
from copy import deepcopy
from pathlib import Path
from typing import Callable

from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution"
INPUTS = ROOT / "environment" / "input_files"
PPT_NAME = "公开课课件_修改完成.pptx"
XLSX_NAME = "修订响应.xlsx"
PROCESSOR_NAME = "revision_requirements_processor.py"
PLAN_NAME = "revision_plan.json"
OUTPUTS = (PPT_NAME, XLSX_NAME, PROCESSOR_NAME, PLAN_NAME)
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS = {"p": P_NS}
Mutator = Callable[[Path], None]


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"无法载入 {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _save_ppt(root: Path, callback) -> None:
    path = root / PPT_NAME
    deck = Presentation(path)
    callback(deck)
    deck.save(path)


def _rewrite_zip(path: Path, member: str, callback) -> None:
    with zipfile.ZipFile(path) as source:
        entries = [(item, source.read(item.filename)) for item in source.infolist()]
    payload = next(data for item, data in entries if item.filename == member)
    root = etree.fromstring(payload)
    callback(root)
    replacement = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(temporary, "w") as target:
        for item, data in entries:
            target.writestr(item, replacement if item.filename == member else data)
    temporary.replace(path)


def _shape(slide, needle: str):
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and needle in shape.text:
            return shape
    raise ValueError(f"未找到文本对象：{needle}")


def _remove_shape(shape) -> None:
    shape._element.getparent().remove(shape._element)


def _plan_change(root: Path, identifier: str, key: str, value) -> None:
    path = root / PLAN_NAME
    plan = json.loads(path.read_text(encoding="utf-8"))
    item = next(entry for entry in plan["requirements"] if entry["id"] == identifier)
    item[key] = value
    path.write_text(json.dumps(plan, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _hardcoded_processor(root: Path) -> None:
    script = """#!/usr/bin/env python3
import argparse
from pathlib import Path
p=argparse.ArgumentParser()
p.add_argument('--requirements')
p.add_argument('--source-pptx')
p.add_argument('--output', type=Path, required=True)
a=p.parse_args()
a.output.write_bytes(Path(__file__).with_name('revision_plan.json').read_bytes())
"""
    (root / PROCESSOR_NAME).write_text(script, encoding="utf-8")


def _cover_font(root: Path) -> None:
    def mutate(deck):
        shape = _shape(deck.slides[0], "Comics & Animation")
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.font.size:
                    run.font.size = run.font.size - 2 * 12700
    _save_ppt(root, mutate)


def _background(root: Path) -> None:
    def mutate(deck):
        background = deck.slides[1]._element.find(".//p:cSld/p:bg", {"p": P_NS})
        color = background.find(".//a:srgbClr", {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"})
        color.set("val", "FFFFFF")
    _save_ppt(root, mutate)


def _slide3_width(root: Path) -> None:
    def mutate(deck):
        picture = next(shape for shape in deck.slides[2].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
        picture.width = int(9 * 360000)
    _save_ppt(root, mutate)


def _slide3_left(root: Path) -> None:
    def mutate(deck):
        picture = next(shape for shape in deck.slides[2].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
        picture.left = int(2 * 360000)
    _save_ppt(root, mutate)


def _slide4_image_right(root: Path) -> None:
    def mutate(deck):
        picture = next(shape for shape in deck.slides[3].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
        picture.left = int(18 * 360000)
    _save_ppt(root, mutate)


def _slide5_remove_image(root: Path) -> None:
    def mutate(deck):
        picture = next(shape for shape in deck.slides[4].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
        _remove_shape(picture)
    _save_ppt(root, mutate)


def _slide6_old_answer(root: Path) -> None:
    def mutate(deck):
        shape = next(shape for shape in deck.slides[5].shapes if getattr(shape, "has_text_frame", False) and "→" in shape.text)
        shape.text = shape.text.split("→", 1)[0] + "→ start the trip"
    _save_ppt(root, mutate)


def _remove_timing(root: Path, page: int) -> None:
    member = f"ppt/slides/slide{page}.xml"
    def mutate(slide):
        timing = slide.find("p:timing", NS)
        if timing is not None:
            timing.getparent().remove(timing)
    _rewrite_zip(root / PPT_NAME, member, mutate)


def _duplicate_animation(root: Path) -> None:
    member = "ppt/slides/slide9.xml"
    a_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    ns = {"p": P_NS, "a": a_ns}
    def mutate(slide):
        effects = slide.xpath(".//p:animEffect", namespaces=ns)
        effects[-1].set("filter", effects[0].get("filter"))
    _rewrite_zip(root / PPT_NAME, member, mutate)


def _slide10_title(root: Path) -> None:
    _save_ppt(root, lambda deck: setattr(_shape(deck.slides[9], "Match five stages"), "text", "Match four stages of the story with frames"))


def _slide10_curve_size(root: Path) -> None:
    def mutate(deck):
        picture = next(shape for shape in deck.slides[9].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
        picture.width = int(15 * 360000)
    _save_ppt(root, mutate)


def _restore_slide16_image(root: Path) -> None:
    source = Presentation(INPUTS / "公开课课件.pptx")
    source_picture = next(shape for shape in source.slides[15].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
    def mutate(deck):
        from io import BytesIO
        deck.slides[15].shapes.add_picture(BytesIO(source_picture.image.blob), source_picture.left, source_picture.top, source_picture.width, source_picture.height)
    _save_ppt(root, mutate)


def _remove_transition(root: Path) -> None:
    def mutate(slide):
        transition = slide.find("p:transition", NS)
        if transition is not None:
            transition.getparent().remove(transition)
    _rewrite_zip(root / PPT_NAME, "ppt/slides/slide1.xml", mutate)


def _duplicate_transition(root: Path) -> None:
    path = root / PPT_NAME
    with zipfile.ZipFile(path) as source:
        entries = [(item, source.read(item.filename)) for item in source.infolist()]
    slide1 = etree.fromstring(next(data for item, data in entries if item.filename == "ppt/slides/slide1.xml"))
    slide2 = etree.fromstring(next(data for item, data in entries if item.filename == "ppt/slides/slide2.xml"))
    first = slide1.find("p:transition", NS)
    second = slide2.find("p:transition", NS)
    if second is not None:
        second.getparent().replace(second, deepcopy(first))
    temporary = path.with_suffix(".pptx.tmp")
    with zipfile.ZipFile(temporary, "w") as target:
        for item, data in entries:
            if item.filename == "ppt/slides/slide2.xml":
                data = etree.tostring(slide2, xml_declaration=True, encoding="UTF-8", standalone=True)
            target.writestr(item, data)
    temporary.replace(path)


def _xlsx_cell(root: Path, identifier: str, column: int, value) -> None:
    path = root / XLSX_NAME
    book = load_workbook(path)
    sheet = book["响应记录"]
    row = next(row for row in range(2, sheet.max_row + 1) if sheet.cell(row, 1).value == identifier)
    sheet.cell(row, column).value = value
    book.save(path)
    book.close()


def _xlsx_omit(root: Path) -> None:
    path = root / XLSX_NAME
    book = load_workbook(path)
    sheet = book["响应记录"]
    row = next(row for row in range(2, sheet.max_row + 1) if sheet.cell(row, 1).value == "R11")
    sheet.delete_rows(row)
    book.save(path)
    book.close()


def _xlsx_extra(root: Path) -> None:
    path = root / XLSX_NAME
    book = load_workbook(path)
    book["响应记录"].append(["R13", "1", "已获得校长批准", "已完成", "pptx:slide=1;check=approved"])
    book.save(path)
    book.close()


def _remove_title(root: Path) -> None:
    def mutate(deck):
        _remove_shape(_shape(deck.slides[11], "Beginning"))
    _save_ppt(root, mutate)


def _text_out_of_bounds(root: Path) -> None:
    def mutate(deck):
        shape = _shape(deck.slides[11], "Beginning")
        shape.left = -int(2 * 360000)
    _save_ppt(root, mutate)


def _unmodified_input(root: Path) -> None:
    shutil.copy2(INPUTS / "公开课课件.pptx", root / PPT_NAME)


def _corrupt(root: Path, name: str) -> None:
    payload = b"def broken(:\n" if name == PROCESSOR_NAME else b"not-a-valid-core-artifact"
    (root / name).write_bytes(payload)


def _moderate(root: Path) -> None:
    _hardcoded_processor(root)
    _slide3_left(root)
    _slide5_remove_image(root)
    _remove_timing(root, 8)
    _slide10_title(root)
    _remove_timing(root, 11)
    _restore_slide16_image(root)
    _remove_transition(root)
    _xlsx_omit(root)


def _severe(root: Path) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Comics & Animation"
    deck.save(root / PPT_NAME)
    path = root / XLSX_NAME
    book = load_workbook(path)
    sheet = book["响应记录"]
    sheet.delete_rows(3, sheet.max_row)
    book.save(path)
    book.close()
    _hardcoded_processor(root)


def _cases() -> list[tuple[str, Mutator, str]]:
    cases: list[tuple[str, Mutator, str]] = []
    for name in OUTPUTS:
        cases.extend([
            (f"missing-{name}", lambda root, item=name: (root / item).unlink(), "fatal_zero"),
            (f"empty-{name}", lambda root, item=name: (root / item).write_bytes(b""), "fatal_zero"),
            (f"corrupt-{name}", lambda root, item=name: _corrupt(root, item), "fatal_zero"),
        ])
    cases.extend([
        ("empty-output", lambda root: [path.unlink() for path in list(root.iterdir()) if path.is_file()], "fatal_zero"),
        ("unmodified-input-renamed", _unmodified_input, "fatal_zero"),
        ("plan-r03-width-wrong", lambda root: _plan_change(root, "R03", "width_cm", 9), "partial"),
        ("plan-r11-page-wrong", lambda root: _plan_change(root, "R11", "slides", [17]), "partial"),
        ("plan-r09-source-text-wrong", lambda root: _plan_change(root, "R09", "source_text", "无来源要求"), "partial"),
        ("hardcoded-processor", _hardcoded_processor, "partial"),
        ("cover-title-font-not-increased", _cover_font, "partial"),
        ("slide2-background-white", _background, "partial"),
        ("slide3-student-image-width-9cm", _slide3_width, "partial"),
        ("slide3-student-image-left", _slide3_left, "partial"),
        ("slide4-adventure-image-right", _slide4_image_right, "partial"),
        ("slide5-one-review-image-missing", _slide5_remove_image, "partial"),
        ("slide6-old-answer-restored", _slide6_old_answer, "partial"),
        ("slide8-all-animation-timing-removed", lambda root: _remove_timing(root, 8), "partial"),
        ("slide9-last-animation-duplicates-first", _duplicate_animation, "partial"),
        ("slide10-title-wrong", _slide10_title, "partial"),
        ("slide10-curve-width-15cm", _slide10_curve_size, "partial"),
        ("slide11-animation-timing-removed", lambda root: _remove_timing(root, 11), "partial"),
        ("slide16-deleted-image-restored", _restore_slide16_image, "partial"),
        ("slide1-transition-missing", _remove_transition, "partial"),
        ("slide2-transition-duplicates-slide1", _duplicate_transition, "partial"),
        ("response-r11-omitted", _xlsx_omit, "partial"),
        ("response-r03-target-conflict", lambda root: _xlsx_cell(root, "R03", 2, "4"), "partial"),
        ("response-r05-source-text-wrong", lambda root: _xlsx_cell(root, "R05", 3, "自行补充图片"), "partial"),
        ("response-r09-status-not-complete", lambda root: _xlsx_cell(root, "R09", 4, "未完成"), "partial"),
        ("response-r10-check-wrong", lambda root: _xlsx_cell(root, "R10", 5, "pptx:slide=11;check=unknown"), "partial"),
        ("response-unsupported-r13", _xlsx_extra, "partial"),
        ("preserved-slide12-title-removed", _remove_title, "partial"),
        ("slide12-title-out-of-bounds", _text_out_of_bounds, "partial"),
        ("extra-debug-file", lambda root: (root / "debug.txt").write_text("debug", encoding="utf-8"), "partial"),
        ("moderate-eight-requirements-broken", _moderate, "moderate"),
        ("severe-one-slide-placeholder", _severe, "severe"),
    ])
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--golden-dir", type=Path, default=GOLDEN)
    parser.add_argument("--input-dir", type=Path, default=INPUTS)
    parser.add_argument("--result", type=Path)
    args = parser.parse_args()
    judge = _load("officeval_080_judge_v2_for_negatives", ROOT / "tests" / "judge.py")
    golden = judge.evaluate(args.golden_dir, args.input_dir, run_render=False)
    if golden.get("reward") != 1.0:
        raise AssertionError(f"Golden v2 必须为 1.0，实际 {golden.get('reward')}")
    results: list[dict] = []
    for name, mutator, layer in _cases():
        with tempfile.TemporaryDirectory(prefix=f"officeval-080-v2-{name}-") as temporary:
            output = Path(temporary) / "output"
            shutil.copytree(args.golden_dir, output)
            mutator(output)
            result = judge.evaluate(output, args.input_dir, run_render=False)
            reward = float(result.get("reward", 0))
            if layer == "fatal_zero" and reward != 0:
                raise AssertionError(f"{name} 应硬归零，实际 {reward}")
            if layer in {"partial", "moderate", "severe"} and not (0 < reward < 1):
                raise AssertionError(f"{name} 应保留连续部分分，实际 {reward}")
            if layer == "moderate" and reward >= 0.8:
                raise AssertionError(f"{name} 应低于通过线，实际 {reward}")
            results.append({"case": name, "layer": layer, "reward": reward, "passed": bool(result.get("passed")), "penalties": result.get("penalties", [])})
    moderate = next(item["reward"] for item in results if item["layer"] == "moderate")
    severe = next(item["reward"] for item in results if item["layer"] == "severe")
    if not 0 < severe < moderate < 0.8:
        raise AssertionError(f"分层应满足 0 < severe({severe}) < moderate({moderate}) < 0.8")
    partials = [item["reward"] for item in results if item["layer"] == "partial"]
    payload = {
        "passed": True,
        "golden_reward": golden["reward"],
        "total": len(results),
        "fatal_zero": sum(item["layer"] == "fatal_zero" for item in results),
        "partial_continuous": sum(item["layer"] != "fatal_zero" for item in results),
        "partial_reward_range": [round(min(partials + [moderate, severe]), 6), round(max(partials + [moderate, severe]), 6)],
        "moderate_reward": moderate,
        "severe_reward": severe,
        "cases": results,
    }
    if args.result:
        args.result.parent.mkdir(parents=True, exist_ok=True)
        args.result.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
