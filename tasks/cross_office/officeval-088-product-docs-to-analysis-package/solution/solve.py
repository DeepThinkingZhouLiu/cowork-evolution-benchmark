#!/usr/bin/env python3
"""逐份阅读并改造 OmegaUse-OfficeVal officeval_088。

脚本只重放这一道真实任务：从五份产品说明 Word 中使用已经人工核对的产品事实，
裁取各自原图中的洗衣机主体，填写上游 Excel 与 Word 模板。它不是通用数据生成器。
"""

from __future__ import annotations

import argparse
import copy
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image


OUTPUT_XLSX = "选货汇总表_已整理.xlsx"
OUTPUT_DOCX = "产品分析.docx"
WORD_TEMPLATE = "模板.docx"
EXCEL_TEMPLATE = "选货汇总表_模板2_已填充 (1).xlsx"

PRODUCTS = [
    {
        "key": "云溪",
        "source": "云溪5.0 D3滚筒洗衣机产品说明文档.docx",
        "name": "云溪5.0 D3滚筒洗衣机",
        "model": "云溪5.0 D3",
        "price": 4500,
        "params": "型号：云溪5.0 D3；洗涤容量：12KG；宽度：595mm；类型：滚筒洗衣机；能效等级：一级；深度：593mm/537mm；高度：850mm",
        "scene": "多人口家庭的日常衣物及大件家纺洗护；小型民宿、公寓出租房",
        "crop": (70, 199, 202, 375),
        "ratio": 0.75,
    },
    {
        "key": "统帅",
        "source": "统帅三桶懒人洗衣机产品说明文档.docx",
        "name": "统帅三桶懒人洗衣机",
        "model": "Leader L10510B12（标尺版）",
        "price": 3999,
        "params": "型号：Leader L10510B12（标尺版）；大筒：10.5KG；双小筒：各1KG；宽度：600mm；深度：570mm；类型：滚筒洗衣机；能效等级：一级",
        "scene": "多人口家庭的日常衣物、大件家纺及外衣/内衣/宝宝衣物分类洗；小型民宿、公寓出租房；母婴家庭除菌衣物清洁",
        "crop": (232, 158, 471, 500),
        "ratio": 0.70,
    },
    {
        "key": "小天鹅",
        "source": "小天鹅12KG滚筒洗衣机产品说明文档.docx",
        "name": "小天鹅12KG滚筒洗衣机",
        "model": "TG12VE40PRO",
        "price": 3039,
        "params": "型号：TG12VE40PRO；洗涤容量：12KG；宽度：595mm；类型：滚筒洗衣机；能效等级：一级；深度：560mm/585mm",
        "scene": "多人口家庭按批次进行外衣、内衣与宝宝衣物分类洗，并可洗护日常衣物和大件家纺；小型民宿、公寓出租房",
        "crop": (105, 367, 359, 669),
        "ratio": 0.84,
    },
    {
        "key": "米家",
        "source": "米家8KG滚筒洗衣机产品说明文档.docx",
        "name": "米家8KG滚筒洗衣机",
        "model": "米家MG80V11D（滚筒款）",
        "price": 999,
        "params": "型号：米家MG80V11D（滚筒款）；洗涤容量：8KG；宽度：598mm；深度：495mm；类型：滚筒洗衣机；能效等级：一级",
        "scene": "单身、2-3口小家庭的日常衣物和薄家纺洗护；小型民宿、员工宿舍",
        "crop": (141, 276, 315, 536),
        "ratio": 0.67,
    },
    {
        "key": "美的",
        "source": "美的8KG波轮洗衣机产品说明文档.docx",
        "name": "美的8KG波轮洗衣机",
        "model": "美的MB80V36DT（波轮款）",
        "price": 846,
        "params": "型号：美的MB80V36DT（波轮款）；洗涤容量：8KG；宽度：515mm；深度：525mm；类型：波轮洗衣机；能效等级：一级",
        "scene": "单身、2-3口小家庭的日常衣物和薄家纺洗护；小型民宿、员工宿舍",
        "crop": (210, 120, 456, 498),
        "ratio": 0.65,
    },
]


def first_media_bytes(path: Path) -> tuple[bytes, str]:
    with zipfile.ZipFile(path) as archive:
        names = [name for name in archive.namelist()
                 if name.startswith("word/media/") and not name.endswith("/")]
        if len(names) != 1:
            raise ValueError(f"{path.name} 的源图片数量为 {len(names)}，期望 1")
        name = names[0]
        return archive.read(name), Path(name).suffix or ".png"


def crop_product_images(input_dir: Path, temp_root: Path) -> dict[str, Path]:
    result = {}
    for product in PRODUCTS:
        raw, suffix = first_media_bytes(input_dir / product["source"])
        source = temp_root / f'{product["key"]}-source{suffix}'
        source.write_bytes(raw)
        cropped = temp_root / f'{product["key"]}.png'
        with Image.open(source) as image:
            image.convert("RGB").crop(product["crop"]).save(cropped, format="PNG")
        with Image.open(cropped) as image:
            ratio = image.width / image.height
            if abs(ratio - product["ratio"]) > 0.012:
                raise ValueError(f'{product["key"]} 裁剪比例 {ratio:.3f} 不符合目标 {product["ratio"]:.2f}')
        result[product["key"]] = cropped
    return result


def style_excel(input_dir: Path, output_dir: Path, images: dict[str, Path]) -> Path:
    workbook = load_workbook(input_dir / EXCEL_TEMPLATE)
    sheet = workbook.active
    headers = ["产品名称", "核心参数", "适用场景", "指导价格", "展示图片"]
    for col, value in enumerate(headers, start=1):
        sheet.cell(1, col).value = value

    widths = {"A": 23.42, "B": 32.81, "C": 10.82, "D": 12.36, "E": 12.08}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="F47B2A")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="微软雅黑", size=11, bold=True, color="000000")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    sheet.row_dimensions[1].height = 24

    for row, product in enumerate(PRODUCTS, start=2):
        values = [product["name"], product["params"], product["scene"], f'{product["price"]}元', None]
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row, col, value)
            cell.font = Font(name="微软雅黑", size=10, color="000000")
            cell.alignment = Alignment(
                horizontal="center" if col in (1, 4, 5) else "left",
                vertical="center",
                wrap_text=True,
            )
            cell.border = border
        sheet.row_dimensions[row].height = 230
        picture = ExcelImage(str(images[product["key"]]))
        picture.height = 230.4
        picture.width = 230.4 * product["ratio"]
        sheet.add_image(picture, f"E{row}")

    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = "A1:E6"
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = "A1:E6"
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.4
    sheet.page_margins.bottom = 0.4

    output_path = output_dir / OUTPUT_XLSX
    workbook.save(output_path)
    return output_path


def split_template_table(document: Document) -> list[Table]:
    original = document.tables[0]._tbl
    parent = original.getparent()
    position = parent.index(original)
    rows = list(original.findall(qn("w:tr")))
    new_elements = []
    for index, row in enumerate(rows):
        table = copy.deepcopy(original)
        for child in list(table.findall(qn("w:tr"))):
            table.remove(child)
        table.append(copy.deepcopy(row))
        parent.insert(position + len(new_elements), table)
        new_elements.append(table)
        if index < len(rows) - 1:
            spacer = OxmlElement("w:p")
            ppr = OxmlElement("w:pPr")
            spacing = OxmlElement("w:spacing")
            spacing.set(qn("w:before"), "20")
            spacing.set(qn("w:after"), "20")
            spacing.set(qn("w:line"), "20")
            spacing.set(qn("w:lineRule"), "exact")
            ppr.append(spacing)
            spacer.append(ppr)
            parent.insert(position + len(new_elements), spacer)
            new_elements.append(spacer)
    parent.remove(original)
    return [Table(element, document._body) for element in new_elements if element.tag == qn("w:tbl")]


def set_cell_margins(cell, value=60):
    tcpr = cell._tc.get_or_add_tcPr()
    margins = tcpr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for edge in ("top", "start", "bottom", "end"):
        item = margins.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            margins.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")


def add_text_line(cell, label: str, value: str, *, title=False):
    paragraph = cell.add_paragraph() if cell.paragraphs[0].text else cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(label)
    run.bold = True
    run.font.name = "微软雅黑"
    run.font.size = Pt(12 if title else 9)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "微软雅黑")
    if value:
        value_run = paragraph.add_run(value)
        value_run.font.name = "微软雅黑"
        value_run.font.size = Pt(9)
        value_run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "微软雅黑")


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_header_line(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    indent = ppr.find(qn("w:ind"))
    if indent is None:
        indent = OxmlElement("w:ind")
        ppr.append(indent)
    indent.set(qn("w:left"), "950")
    indent.set(qn("w:right"), "950")
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)
    ppr.append(borders)


def style_word(input_dir: Path, output_dir: Path, images: dict[str, Path]) -> Path:
    document = Document(input_dir / WORD_TEMPLATE)
    tables = split_template_table(document)
    selected = PRODUCTS[:3]
    for index, (table, product) in enumerate(zip(tables, selected), start=1):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        row = table.rows[0]
        row.height = Cm(5.1)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        trpr = row._tr.get_or_add_trPr()
        if trpr.find(qn("w:cantSplit")) is None:
            trpr.append(OxmlElement("w:cantSplit"))
        left, right = row.cells
        for cell in (left, right):
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
        add_text_line(left, f"产品{['一', '二', '三'][index - 1]}：", "", title=True)
        add_text_line(left, "产品名称：", product["name"])
        add_text_line(left, "核心参数：", product["params"])
        add_text_line(left, "适用场景：", product["scene"])
        add_text_line(left, "指导价格：", f'{product["price"]}元')
        picture_paragraph = right.paragraphs[0]
        picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture_paragraph.paragraph_format.space_before = Pt(0)
        picture_paragraph.paragraph_format.space_after = Pt(0)
        picture_paragraph.add_run().add_picture(
            str(images[product["key"]]),
            height=Cm(5.0),
            width=Cm(5.0 * product["ratio"]),
        )

    section = document.sections[0]
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.text = "产品分析"
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_paragraph.paragraph_format.space_after = Pt(3)
    for run in header_paragraph.runs:
        run.font.name = "微软雅黑"
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "微软雅黑")
    add_header_line(header_paragraph)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.text = ""
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(footer_paragraph)

    output_path = output_dir / OUTPUT_DOCX
    document.save(output_path)
    return output_path


def create_outputs(input_dir: Path, output_dir: Path):
    for product in PRODUCTS:
        if not (input_dir / product["source"]).is_file():
            raise FileNotFoundError(input_dir / product["source"])
    output_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="officeval-088-") as temp_dir:
        images = crop_product_images(input_dir, Path(temp_dir))
        xlsx = style_excel(input_dir, output_dir, images)
        docx = style_word(input_dir, output_dir, images)
    return xlsx, docx


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    print(*create_outputs(args.input_dir, args.output_dir), sep="\n")
