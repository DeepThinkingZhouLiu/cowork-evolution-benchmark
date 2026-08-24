#!/usr/bin/env python3
"""officeval_088 的 Harbor Judge：真实源事实、来源裁图与双输出一致性。"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from openpyxl import load_workbook
from PIL import Image, ImageChops, ImageStat


TASK_ID = "cowork-evo/officeval-088-product-docs-to-analysis-package"
RUBRIC_VERSION = "omegause-officeval-088-adaptation-v0.1"
JUDGE_VERSION = "officeval-088-adaptation-v0.1"
OUTPUT_XLSX = "选货汇总表_已整理.xlsx"
OUTPUT_DOCX = "产品分析.docx"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
NS = {"w": W_NS, "wp": WP_NS}

CROPS = {
    "云溪": ("云溪5.0 D3滚筒洗衣机产品说明文档.docx", (70, 199, 202, 375)),
    "统帅": ("统帅三桶懒人洗衣机产品说明文档.docx", (232, 158, 471, 500)),
    "小天鹅": ("小天鹅12KG滚筒洗衣机产品说明文档.docx", (105, 367, 359, 669)),
    "米家": ("米家8KG滚筒洗衣机产品说明文档.docx", (141, 276, 315, 536)),
    "美的": ("美的8KG波轮洗衣机产品说明文档.docx", (210, 120, 456, 498)),
}

CRITERION_SPECS = {
    "R001": ("指定 XLSX 与 DOCX 均存在、可打开且可编辑", 3, "hurdle"),
    "R002": ("Excel 五个产品的名称、型号、价格和核心参数准确", 6, "positive"),
    "R003": ("Excel 适用场景来自对应源说明且长文本不溢出", 3, "positive"),
    "R004": ("Excel 五张图片来自对应源图裁剪并正确锚定", 4, "positive"),
    "R005": ("Excel 表头、列宽、边框、换行和打印布局满足模板", 3, "positive"),
    "R006": ("Word 高价前三名顺序、事实与 Excel 一致", 5, "positive"),
    "R007": ("Word 三张来源裁图、尺寸和产品块不可跨页正确", 3, "positive"),
    "R008": ("Word 页眉页码与两份文件实际渲染完整", 3, "positive"),
    "R009": ("价格错误或两个以上产品关键事实错误", -8, "penalty"),
    "R010": ("两张以上产品图片缺失、错行或不是来源裁图", -6, "penalty"),
    "R011": ("Word 高价顺序错误或双输出两处以上不一致", -8, "penalty"),
    "R012": ("渲染失败、关键行裁切、产品跨页或页眉页码缺失", -5, "penalty"),
}


def norm(text: str) -> str:
    return re.sub(r"\s+", "", str(text or ""))


def criterion(identifier: str, passed: bool, evidence: str = "") -> dict:
    description, weight, kind = CRITERION_SPECS[identifier]
    return {
        "criterion_id": identifier,
        "description": description,
        "weight": weight,
        "type": kind,
        "passed": bool(passed),
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
    }


def expected_crops(input_dir: Path) -> dict[str, Image.Image]:
    result = {}
    for key, (filename, box) in CROPS.items():
        with zipfile.ZipFile(input_dir / filename) as archive:
            media = [name for name in archive.namelist()
                     if name.startswith("word/media/") and not name.endswith("/")]
            if len(media) != 1:
                continue
            with Image.open(io.BytesIO(archive.read(media[0]))) as image:
                result[key] = image.convert("RGB").crop(box)
    return result


def image_distance(actual_bytes: bytes, expected: Image.Image) -> float:
    try:
        with Image.open(io.BytesIO(actual_bytes)) as image:
            actual = image.convert("RGB").resize((64, 64))
        reference = expected.convert("RGB").resize((64, 64))
        difference = ImageChops.difference(actual, reference)
        return sum(ImageStat.Stat(difference).mean) / 3.0
    except Exception:
        return 255.0


def excel_state(path: Path, expectations: dict, crops: dict[str, Image.Image]) -> dict:
    workbook = load_workbook(path, data_only=False)
    sheet = workbook.active
    rows = {}
    for row in range(2, sheet.max_row + 1):
        name = norm(sheet.cell(row, 1).value)
        for product in expectations["products"]:
            if norm(product["name"]) in name:
                rows[product["key"]] = row
    fact_issues = []
    scene_issues = []
    price_errors = []
    product_errors = 0
    for product in expectations["products"]:
        key = product["key"]
        row = rows.get(key)
        if not row:
            fact_issues.append(f"{key}:缺行")
            product_errors += 1
            continue
        joined = norm(" ".join(str(sheet.cell(row, col).value or "") for col in range(1, 5)))
        missing = [item for item in [product["name"], product["model"], *product["parameters"]]
                   if norm(item) not in joined]
        actual_price = norm(sheet.cell(row, 4).value)
        if str(product["price"]) not in re.sub(r"\D", "", actual_price):
            price_errors.append(f"{key}:{actual_price}")
        if missing:
            fact_issues.append(f"{key}:缺{missing}")
            product_errors += 1
        scene = norm(sheet.cell(row, 3).value)
        missing_scenes = [item for item in product["scenes"] if norm(item) not in scene]
        if missing_scenes:
            scene_issues.append(f"{key}:缺{missing_scenes}")
        if not sheet.cell(row, 2).alignment.wrap_text or not sheet.cell(row, 3).alignment.wrap_text:
            scene_issues.append(f"{key}:核心参数或场景未自动换行")
        if (sheet.row_dimensions[row].height or 15) < 170:
            scene_issues.append(f"{key}:行高={sheet.row_dimensions[row].height}")

    image_issues = []
    image_details = []
    for image in sheet._images:
        anchor = getattr(image.anchor, "_from", None)
        if anchor is None:
            image_issues.append("存在无法定位锚点的图片")
            continue
        row, col = anchor.row + 1, anchor.col + 1
        product = next((p for p in expectations["products"] if rows.get(p["key"]) == row), None)
        if product is None or col != 5:
            image_issues.append(f"图片锚点=({row},{col})")
            continue
        key = product["key"]
        raw = image._data()
        distance = image_distance(raw, crops[key])
        extent = getattr(image.anchor, "ext", None)
        cx = int(getattr(extent, "cx", 0) or 0)
        cy = int(getattr(extent, "cy", 0) or 0)
        ratio = cx / cy if cy else 0
        height_chars = cy / 182880 if cy else 0
        if distance > 35 or abs(ratio - product["ratio"]) > 0.02 or not 11.8 <= height_chars <= 12.1:
            image_issues.append(f"{key}:差异={distance:.1f},比={ratio:.3f},高={height_chars:.2f}")
        image_details.append(f"{key}:差异={distance:.1f},比={ratio:.3f},高={height_chars:.2f}")
    if len(sheet._images) != 5:
        image_issues.append(f"图片数={len(sheet._images)}")

    expected_widths = [23.42, 32.81, 10.82, 12.36, 12.08]
    width_issues = []
    for index, expected in enumerate(expected_widths, start=1):
        actual = sheet.column_dimensions[chr(64 + index)].width or 8.43
        if abs(actual - expected) > 0.1:
            width_issues.append(f"{chr(64+index)}={actual:.2f}")
    headers = [norm(sheet.cell(1, col).value) for col in range(1, 6)]
    style_ok = headers == ["产品名称", "核心参数", "适用场景", "指导价格", "展示图片"]
    style_ok = style_ok and not width_issues and sheet.page_setup.orientation == "landscape"
    return {
        "workbook": workbook,
        "sheet": sheet,
        "rows": rows,
        "fact_issues": fact_issues,
        "scene_issues": scene_issues,
        "price_errors": price_errors,
        "product_errors": product_errors,
        "image_issues": image_issues,
        "image_details": image_details,
        "image_count": len(sheet._images),
        "style_ok": style_ok,
        "style_evidence": f"表头={headers}，列宽问题={width_issues}，方向={sheet.page_setup.orientation}",
    }


def table_text(table) -> str:
    return norm(" ".join(cell.text for row in table.rows for cell in row.cells))


def word_image_for_table(document: Document, table) -> tuple[bytes, float, float] | None:
    blips = table._tbl.xpath(".//a:blip")
    extents = table._tbl.xpath(".//wp:extent")
    if not blips or not extents:
        return None
    rid = blips[0].get(qn("r:embed"))
    part = document.part.related_parts.get(rid)
    if part is None:
        return None
    cx = int(extents[0].get("cx", "0"))
    cy = int(extents[0].get("cy", "0"))
    return part.blob, cx / cy if cy else 0, cy / 360000 if cy else 0


def word_state(path: Path, expectations: dict, crops: dict[str, Image.Image], excel: dict) -> dict:
    document = Document(path)
    blocks = []
    for table in document.tables:
        text = table_text(table)
        key = next((p["key"] for p in expectations["products"] if norm(p["name"]) in text), None)
        if key:
            blocks.append((key, text, table))
    expected_order = expectations["top_three"]
    actual_order = [item[0] for item in blocks]
    order_ok = actual_order == expected_order
    fact_issues = []
    inconsistencies = []
    image_issues = []
    image_details = []
    product_by_key = {p["key"]: p for p in expectations["products"]}
    for key, text, table in blocks:
        product = product_by_key[key]
        required = [product["name"], product["model"], str(product["price"]),
                    *product["parameters"], *product["scenes"]]
        missing = [item for item in required if norm(item) not in text]
        if missing:
            fact_issues.append(f"{key}:缺{missing}")
        row = excel["rows"].get(key)
        if row:
            excel_text = norm(" ".join(str(excel["sheet"].cell(row, col).value or "") for col in range(1, 5)))
            for token in [product["name"], product["model"], str(product["price"]), *product["parameters"]]:
                if norm(token) not in text or norm(token) not in excel_text:
                    inconsistencies.append(f"{key}:{token}")
        picture = word_image_for_table(document, table)
        if picture is None:
            image_issues.append(f"{key}:无图")
        else:
            raw, ratio, height_cm = picture
            distance = image_distance(raw, crops[key])
            if distance > 35 or abs(ratio - product["ratio"]) > 0.02 or not 4.9 <= height_cm <= 5.1:
                image_issues.append(f"{key}:差异={distance:.1f},比={ratio:.3f},高={height_cm:.2f}")
            image_details.append(f"{key}:差异={distance:.1f},比={ratio:.3f},高={height_cm:.2f}")
        row_el = table.rows[0]._tr
        trpr = row_el.trPr
        if trpr is None or trpr.find(qn("w:cantSplit")) is None:
            image_issues.append(f"{key}:产品块可跨页")

    with zipfile.ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml")
        header_xml = b"".join(archive.read(name) for name in archive.namelist()
                              if name.startswith("word/header") and name.endswith(".xml"))
        footer_xml = b"".join(archive.read(name) for name in archive.namelist()
                              if name.startswith("word/footer") and name.endswith(".xml"))
    header_text = header_xml.decode("utf-8", errors="ignore")
    footer_text = footer_xml.decode("utf-8", errors="ignore")
    header_ok = "产品分析" in header_text and "pBdr" in header_text and "bottom" in header_text
    footer_ok = "PAGE" in footer_text and 'w:val="center"' in footer_text
    return {
        "document": document,
        "blocks": blocks,
        "actual_order": actual_order,
        "order_ok": order_ok,
        "fact_issues": fact_issues,
        "inconsistencies": inconsistencies,
        "image_issues": image_issues,
        "image_details": image_details,
        "header_ok": header_ok,
        "footer_ok": footer_ok,
        "document_xml_bytes": len(document_xml),
    }


def render_package(xlsx_path: Path, docx_path: Path) -> tuple[bool, str]:
    libreoffice = shutil.which("libreoffice")
    pdftotext = shutil.which("pdftotext")
    pdfinfo = shutil.which("pdfinfo")
    if not libreoffice or not pdftotext or not pdfinfo:
        return False, "缺少 libreoffice、pdftotext 或 pdfinfo"
    evidence = []
    with tempfile.TemporaryDirectory(prefix="cowork-evo-088-") as temp_dir:
        root = Path(temp_dir)
        for source, expected_pages, required in [
            (docx_path, 1, ["产品摘要", "云溪", "统帅", "小天鹅", "4500", "3999", "3039"]),
            (xlsx_path, 3, ["云溪", "统帅", "小天鹅", "米家", "美的", "4500", "3999", "3039", "999", "846"]),
        ]:
            profile = (root / f"profile-{source.suffix[1:]}").resolve().as_uri()
            run = subprocess.run(
                [libreoffice, "--headless", f"-env:UserInstallation={profile}",
                 "--convert-to", "pdf", "--outdir", temp_dir, str(source)],
                capture_output=True, text=True, timeout=300, check=False,
            )
            pdf = root / f"{source.stem}.pdf"
            if run.returncode != 0 or not pdf.is_file():
                return False, f"{source.name} 转 PDF 失败：{(run.stderr or run.stdout)[-500:]}"
            info = subprocess.run([pdfinfo, str(pdf)], capture_output=True, text=True, timeout=120)
            match = re.search(r"^Pages:\s+(\d+)", info.stdout, re.MULTILINE)
            pages = int(match.group(1)) if match else 0
            extracted = subprocess.run([pdftotext, "-layout", str(pdf), "-"],
                                       capture_output=True, text=True, timeout=120)
            rendered = norm(extracted.stdout)
            missing = [item for item in required if norm(item) not in rendered]
            evidence.append(f"{source.name}:页={pages},缺失={missing}")
            if pages != expected_pages or missing:
                return False, "；".join(evidence)
    return True, "；".join(evidence)


def evaluate(artifact_dir: Path) -> dict:
    xlsx_path = artifact_dir / OUTPUT_XLSX
    docx_path = artifact_dir / OUTPUT_DOCX
    expectations = json.loads(Path(__file__).with_name("expectations.json").read_text(encoding="utf-8"))
    input_dir = Path(__file__).resolve().parents[1] / "data" / "input_files"
    results = []
    if not xlsx_path.is_file() or not docx_path.is_file():
        results.append(criterion("R001", False,
                                 f"XLSX={xlsx_path.is_file()}，DOCX={docx_path.is_file()}"))
        return build_result(results, [xlsx_path, docx_path])
    try:
        crops = expected_crops(input_dir)
        excel = excel_state(xlsx_path, expectations, crops)
        word = word_state(docx_path, expectations, crops, excel)
    except Exception as exc:
        results.append(criterion("R001", False, f"文件无法解析：{type(exc).__name__}: {exc}"))
        return build_result(results, [xlsx_path, docx_path])

    results.append(criterion("R001", True, "指定 XLSX 与 DOCX 均可解析为可编辑 OOXML"))
    r2_ok = not excel["fact_issues"] and not excel["price_errors"] and len(excel["rows"]) == 5
    results.append(criterion("R002", r2_ok,
                             f"行={sorted(excel['rows'])}，事实问题={excel['fact_issues']}，价格问题={excel['price_errors']}"))
    r3_ok = not excel["scene_issues"]
    results.append(criterion("R003", r3_ok, f"场景/换行问题={excel['scene_issues']}"))
    r4_ok = not excel["image_issues"]
    results.append(criterion("R004", r4_ok,
                             f"图片问题={excel['image_issues']}；{'；'.join(excel['image_details'])}"))
    results.append(criterion("R005", excel["style_ok"], excel["style_evidence"]))

    r6_ok = word["order_ok"] and not word["fact_issues"] and not word["inconsistencies"]
    results.append(criterion("R006", r6_ok,
                             f"顺序={word['actual_order']}，事实问题={word['fact_issues']}，不一致={word['inconsistencies']}"))
    r7_ok = len(word["blocks"]) == 3 and not word["image_issues"]
    results.append(criterion("R007", r7_ok,
                             f"产品块={len(word['blocks'])}，图片问题={word['image_issues']}；{'；'.join(word['image_details'])}"))
    rendered_ok, render_evidence = render_package(xlsx_path, docx_path)
    r8_ok = word["header_ok"] and word["footer_ok"] and rendered_ok
    results.append(criterion("R008", r8_ok,
                             f"页眉={word['header_ok']}，页脚={word['footer_ok']}；{render_evidence}"))

    results.append(criterion("R009", bool(excel["price_errors"]) or excel["product_errors"] >= 2,
                             f"价格错误={excel['price_errors']}，产品错误数={excel['product_errors']}"))
    missing_image_count = max(0, 5 - excel["image_count"])
    results.append(criterion("R010", missing_image_count >= 2 or len(excel["image_issues"]) >= 2,
                             f"图片数={excel['image_count']}，缺失数={missing_image_count}，问题={excel['image_issues']}"))
    results.append(criterion("R011", not word["order_ok"] or len(word["inconsistencies"]) >= 2,
                             f"顺序={word['actual_order']}，不一致={word['inconsistencies']}"))
    results.append(criterion("R012", not r8_ok,
                             f"页眉={word['header_ok']}，页脚={word['footer_ok']}；{render_evidence}"))
    return build_result(results, [xlsx_path, docx_path])


def build_result(results: list[dict], artifact_paths: list[Path]) -> dict:
    present = {item["criterion_id"] for item in results}
    for identifier in CRITERION_SPECS:
        if identifier not in present:
            results.append(criterion(identifier, False, "因交付不存在、损坏或门槛失败，本判据未执行"))
    results.sort(key=lambda item: item["criterion_id"])
    positive_total = sum(item["weight"] for item in results if item["weight"] > 0)
    earned = sum(item["weight"] for item in results if item["weight"] > 0 and item["passed"])
    penalties = sum(item["weight"] for item in results if item["weight"] < 0 and item["passed"])
    reward = max(0.0, min(1.0, (earned + penalties) / positive_total))
    hurdle_failures = [item["criterion_id"] for item in results
                       if item["type"] == "hurdle" and not item["passed"]]
    if hurdle_failures:
        reward = 0.0
    digests = [hashlib.sha256(path.read_bytes()).hexdigest() for path in artifact_paths if path.is_file()]
    if not digests:
        digests = [hashlib.sha256(b"").hexdigest()]
    normalized_results = []
    for item in results:
        score = (0.0 if item["passed"] else 1.0) if item["type"] == "penalty" else item["score"]
        normalized_results.append({
            "criterion_id": item["criterion_id"],
            "score": score,
            "evidence": item["evidence"],
            "method": item["method"],
            "raw": {"description": item["description"], "weight": item["weight"],
                    "type": item["type"], "passed": item["passed"]},
        })
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": digests,
        "criterion_results": normalized_results,
        "reward": round(reward, 6),
        "passed": reward >= 0.75,
        "hurdle_failures": hurdle_failures,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path)
    args = parser.parse_args()
    result = evaluate(args.artifact_dir)
    rendered = json.dumps(result, ensure_ascii=False, indent=2)
    print(rendered)
    if args.result:
        args.result.parent.mkdir(parents=True, exist_ok=True)
        args.result.write_text(rendered + "\n", encoding="utf-8")
    raise SystemExit(0 if result["passed"] else 1)


if __name__ == "__main__":
    main()
