#!/usr/bin/env python3
"""OfficeVal 088 v2 的核心失败与原子部分完成反例。"""

from __future__ import annotations

import importlib.util
import json
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Callable

from docx import Document
from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "solution"
XLSX = "选货汇总表_已整理.xlsx"
DOCX = "产品分析.docx"
Mutator = Callable[[Path], None]


def _load_judge():
    path = ROOT / "tests" / "judge.py"
    spec = importlib.util.spec_from_file_location("officeval_088_v2_judge", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法载入本题 Judge")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


JUDGE = _load_judge()


def _copy(destination: Path) -> None:
    shutil.copytree(GOLDEN, destination)


def _excel_cell(root: Path, coordinate: str, value: object) -> None:
    path = root / XLSX
    book = load_workbook(path)
    book.active[coordinate] = value
    book.save(path)
    book.close()


def _excel_remove_one_image(root: Path) -> None:
    path = root / XLSX
    book = load_workbook(path)
    sheet = book.active
    sheet._images = list(sheet._images[1:])
    book.save(path)
    book.close()


def _excel_width(root: Path) -> None:
    path = root / XLSX
    book = load_workbook(path)
    book.active.column_dimensions["A"].width = 12
    book.save(path)
    book.close()


def _word_replace(root: Path, old: str, new: str) -> None:
    path = root / DOCX
    document = Document(path)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if old in paragraph.text:
                        for run in paragraph.runs:
                            if old in run.text:
                                run.text = run.text.replace(old, new, 1)
                                document.save(path)
                                return
    raise ValueError(f"Word 中未找到 {old}")


def _word_remove_one_image(root: Path) -> None:
    path = root / DOCX
    document = Document(path)
    drawing = document.tables[0]._tbl.xpath(".//w:drawing")[0]
    drawing.getparent().remove(drawing)
    document.save(path)


def _word_clear_header(root: Path) -> None:
    path = root / DOCX
    document = Document(path)
    for section in document.sections:
        for paragraph in section.header.paragraphs:
            paragraph.text = ""
    document.save(path)


def _word_clear_footer(root: Path) -> None:
    path = root / DOCX
    document = Document(path)
    for section in document.sections:
        for paragraph in section.footer.paragraphs:
            paragraph.text = ""
    document.save(path)


def _severe_valid_incomplete(root: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.append(["产品名称", "核心参数", "适用场景", "指导价格", "展示图片"])
    book.save(root / XLSX)
    book.close()
    document = Document()
    document.add_heading("产品分析", level=1)
    document.add_paragraph("产品信息尚未整理完成。")
    document.save(root / DOCX)


def _cases() -> list[tuple[str, Mutator, str]]:
    rows: list[tuple[str, Mutator, str]] = []
    for filename in (XLSX, DOCX):
        rows.extend(
            [
                (f"missing-{filename}", lambda root, name=filename: (root / name).unlink(), "fatal_zero"),
                (f"empty-{filename}", lambda root, name=filename: (root / name).write_bytes(b""), "fatal_zero"),
                (f"corrupt-{filename}", lambda root, name=filename: (root / name).write_bytes(b"not-an-office-file"), "fatal_zero"),
            ]
        )
    rows.extend(
        [
            ("excel-one-price-wrong", lambda root: _excel_cell(root, "D2", 4501), "partial"),
            ("excel-one-scene-incomplete", lambda root: _excel_cell(root, "C2", "多人口家庭"), "partial"),
            ("excel-one-image-missing", _excel_remove_one_image, "partial"),
            ("excel-one-column-width-wrong", _excel_width, "partial"),
            ("word-one-price-wrong", lambda root: _word_replace(root, "4500", "4501"), "partial"),
            ("word-one-image-missing", _word_remove_one_image, "partial"),
            ("word-header-missing", _word_clear_header, "partial"),
            ("word-footer-missing", _word_clear_footer, "partial"),
            ("severe-valid-incomplete", _severe_valid_incomplete, "severe"),
        ]
    )
    return rows


def main() -> None:
    golden = JUDGE.evaluate(GOLDEN, run_render=True)
    if golden["reward"] != 1.0:
        raise AssertionError(f"Golden v2 必须为 1.0，实际 {golden['reward']}")
    records: list[dict] = []
    with tempfile.TemporaryDirectory(prefix="officeval-088-v2-negatives-") as temporary:
        base = Path(temporary)
        empty = base / "fatal-empty-output"
        empty.mkdir()
        reward = JUDGE.evaluate(empty, run_render=False)["reward"]
        records.append({"case": "empty-output", "layer": "fatal_zero", "reward": reward})
        if reward != 0:
            raise AssertionError(f"空交付应为 0，实际 {reward}")

        for name, mutate, layer in _cases():
            case = base / name
            _copy(case)
            mutate(case)
            reward = float(JUDGE.evaluate(case, run_render=False)["reward"])
            if layer == "fatal_zero" and reward != 0:
                raise AssertionError(f"{name} 应为 0，实际 {reward}")
            if layer == "partial" and not 0 < reward < 1:
                raise AssertionError(f"{name} 应保留连续部分分，实际 {reward}")
            if layer == "severe" and not 0 <= reward < 0.5:
                raise AssertionError(f"{name} 应明显低于 0.5，实际 {reward}")
            records.append({"case": name, "layer": layer, "reward": reward})

    partial = [row["reward"] for row in records if row["layer"] == "partial"]
    report = {
        "golden_reward": golden["reward"],
        "total": len(records),
        "fatal_zero": sum(row["layer"] == "fatal_zero" and row["reward"] == 0 for row in records),
        "partial_continuous": len(partial),
        "partial_reward_range": [min(partial), max(partial)],
        "cases": records,
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
