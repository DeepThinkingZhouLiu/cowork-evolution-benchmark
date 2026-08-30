#!/usr/bin/env python3
"""AutomationBench 1115 逐源综合的 StrataCorp v4 竞争响应简报 Golden。"""

from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter
from email import policy
from email.parser import BytesParser
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


HASHES = {
    "Competitive_Intelligence_Tracker.xlsx": "4bb463f0b76e21a875586f2f30ce666cde0a98280fa1d4ab281fa6d19c4f8883",
    "Analyst_Note_StrataCorp_v4.eml": "58854e7344764bef72b7e4b0ae1fd9313b3ad3308293792292da2d50f445bdee",
    "Pricing_Page_Change_Policy.eml": "36fd563f10d5a777d7c6ca90db270961011a3bab3481dea849ace4c69135dd7d",
    "Competitive_Intel_Slack.json": "f9fc7f315c4c4ba7fddf18dd4a64f8225846a45fa146921cad50938624d8b665",
    "AutomationBench_1115_Context.json": "6969d63a0d6163bbdaca24ac0495b73f3e0875997f54bf5143564fe3ab0102ed",
}

NAVY = RGBColor(0x0B, 0x1F, 0x33)
BLUE = RGBColor(0x17, 0x72, 0xD0)
CYAN = RGBColor(0x22, 0xB8, 0xCF)
GREEN = RGBColor(0x17, 0xA6, 0x73)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
RED = RGBColor(0xD9, 0x2D, 0x20)
PURPLE = RGBColor(0x69, 0x47, 0xC6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1D, 0x29, 0x39)
GRAY = RGBColor(0x47, 0x55, 0x69)
MUTED = RGBColor(0x66, 0x70, 0x85)
LINE = RGBColor(0xE4, 0xE7, 0xEC)
SURFACE = RGBColor(0xF7, 0xF9, 0xFC)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def add_text(slide, text, left, top, width, height, *, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = shape.text_frame; frame.clear(); frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]; paragraph.alignment = align
    run = paragraph.add_run(); run.text = str(text); run.font.name = "Aptos"; run.font.size = Pt(size); run.font.color.rgb = color; run.font.bold = bold
    return shape


def add_rect(slide, left, top, width, height, fill, *, line=None, rounded=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill; shape.line.color.rgb = line or fill
    return shape


def background(slide, color):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = color


def header(slide, title, kicker, page):
    add_text(slide, kicker.upper(), 0.65, 0.28, 7.5, 0.25, size=10, color=BLUE, bold=True)
    add_text(slide, title, 0.65, 0.58, 11.8, 0.55, size=27, color=NAVY, bold=True)
    add_text(slide, f"{page:02d} / 06", 11.7, 0.34, 0.9, 0.22, size=9, color=MUTED, align=PP_ALIGN.RIGHT)
    add_rect(slide, 0.65, 1.2, 12.03, 0.02, LINE, rounded=False)


def footer(slide, page, dark=False):
    color = RGBColor(0xA9, 0xB8, 0xC9) if dark else MUTED
    add_text(slide, "AutomationBench 1115  ·  Multi-source snapshot 2026-01-27", 0.65, 7.14, 8, 0.18, size=8, color=color)
    add_text(slide, str(page), 12.2, 7.14, 0.4, 0.18, size=8, color=color, align=PP_ALIGN.RIGHT)


def pill(slide, text, left, top, width, fill, color=WHITE):
    add_rect(slide, left, top, width, 0.34, fill)
    add_text(slide, text, left, top + 0.045, width, 0.2, size=9, color=color, bold=True, align=PP_ALIGN.CENTER)


def metric(slide, value, label, left, top, width, accent):
    add_rect(slide, left, top, width, 1.1, WHITE, line=LINE)
    add_rect(slide, left, top, 0.07, 1.1, accent, rounded=False)
    add_text(slide, value, left + 0.22, top + 0.16, width - 0.3, 0.42, size=24, color=NAVY, bold=True)
    add_text(slide, label.upper(), left + 0.22, top + 0.7, width - 0.3, 0.22, size=9, color=MUTED, bold=True)


def read_inputs(root: Path):
    for name, expected in HASHES.items():
        path = root / name
        if not path.is_file() or sha256(path) != expected:
            raise ValueError(f"冻结输入 SHA256 不匹配：{name}")
    wb = load_workbook(root / "Competitive_Intelligence_Tracker.xlsx", data_only=True)
    ws = wb["Feature Comparison"]
    headers = [cell.value for cell in ws[1]]
    rows = [dict(zip(headers, [cell.value for cell in row], strict=True)) for row in ws.iter_rows(min_row=2)]
    analyst = BytesParser(policy=policy.default).parsebytes((root / "Analyst_Note_StrataCorp_v4.eml").read_bytes()).get_body(preferencelist=("plain",)).get_content()
    pricing_policy = BytesParser(policy=policy.default).parsebytes((root / "Pricing_Page_Change_Policy.eml").read_bytes()).get_body(preferencelist=("plain",)).get_content()
    slack = json.loads((root / "Competitive_Intel_Slack.json").read_text())
    context = json.loads((root / "AutomationBench_1115_Context.json").read_text())
    if len(rows) != 6 or context["deliver_to"] != "strategy@company.example.com":
        raise ValueError("上游 feature 行数或交付受众不匹配")
    return rows, analyst, pricing_policy, slack, context


def build(input_dir: Path, output: Path) -> None:
    rows, analyst, pricing_policy, slack, context = read_inputs(input_dir)
    status_counts = Counter(row["gap_status"] for row in rows)
    if status_counts != Counter({"Parity": 2, "Behind": 2, "Ahead": 1, "Gap": 1}):
        raise ValueError("已审计 gap status 分布不匹配")
    required_analyst = ["3x", "4.2/5", "$29/user/mo", "$79/user/mo", "15-20%", "12 trigger types", "30+"]
    required_slack = ["Meridian Corp", "30%", "~$45K ARR", "Two prospects", "NovaTech", "basically nonexistent"]
    if not all(token in analyst for token in required_analyst) or not all(any(token in message["text"] for message in slack["messages"]) for token in required_slack):
        raise ValueError("上游邮件或 Slack 关键事实不匹配")
    if "Do NOT make direct changes" not in pricing_policy:
        raise ValueError("价格治理政策不匹配")

    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5); blank = prs.slide_layouts[6]

    # 1. Cover
    slide = prs.slides.add_slide(blank); background(slide, NAVY)
    add_text(slide, "COMPETITIVE RESPONSE  /  27 JAN 2026", 0.72, 0.64, 7.5, 0.25, size=10, color=CYAN, bold=True)
    add_text(slide, "StrataCorp v4\nCompetitive Response Brief", 0.72, 1.15, 7.8, 1.55, size=34, color=WHITE, bold=True)
    add_text(slide, "Tracker + analyst note + field intelligence", 0.75, 2.95, 7.2, 0.34, size=14, color=RGBColor(0xCF, 0xD8, 0xE3))
    add_rect(slide, 9.06, 0.95, 3.5, 2.5, RGBColor(0x13, 0x36, 0x52), line=RGBColor(0x2A, 0x53, 0x72))
    add_text(slide, "MIXED", 9.42, 1.38, 2.8, 0.56, size=30, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "COMPETITIVE POSTURE", 9.42, 2.08, 2.8, 0.25, size=10, color=RGBColor(0xA9, 0xC7, 0xDC), bold=True, align=PP_ALIGN.CENTER)
    pill(slide, "ACT NOW", 9.85, 2.62, 1.92, RED)
    metric(slide, "6", "tracked capabilities", 0.72, 4.85, 3.75, BLUE)
    metric(slide, "3", "behind / gap", 4.79, 4.85, 3.75, RED)
    metric(slide, "1", "clear advantage", 8.86, 4.85, 3.75, GREEN)
    add_text(slide, "Delivery: strategy@company.example.com", 0.75, 6.4, 6.2, 0.28, size=11, color=RGBColor(0xCF, 0xD8, 0xE3))
    footer(slide, 1, dark=True)

    # 2. Portfolio posture
    slide = prs.slides.add_slide(blank); background(slide, WHITE); header(slide, "Where we stand", "Portfolio posture", 2)
    data = ChartData(); data.categories = ["Parity", "Behind", "Ahead", "Gap"]; data.add_series("Capabilities", [2, 2, 1, 1])
    shape = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.72), Inches(1.52), Inches(6.0), Inches(4.95), data)
    chart = shape.chart; chart.has_legend = False; chart.has_title = False; chart.value_axis.maximum_scale = 2.5; chart.value_axis.minimum_scale = 0
    plot = chart.plots[0]; plot.has_data_labels = True; plot.data_labels.show_value = True; plot.data_labels.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
    plot.series[0].format.fill.solid(); plot.series[0].format.fill.fore_color.rgb = BLUE
    add_text(slide, "Headline moves", 7.2, 1.58, 4.9, 0.35, size=18, color=NAVY, bold=True)
    cards = [
        ("AI-Powered Insights", "P0 Behind", "Native GPT-4 integration; 3x faster reporting", RED),
        ("Native Mobile App", "P1 Gap", "New v4 app; beta users rate it 4.2/5", AMBER),
        ("Embedded Analytics", "P1 Behind", "DataViz Co. acquisition now GA", RED),
        ("Custom Workflows", "P1 Ahead", "Our strongest counter-position", GREEN),
    ]
    for index, (title, label, detail, accent) in enumerate(cards):
        top = 2.08 + index * 1.0
        add_rect(slide, 7.2, top, 5.2, 0.83, SURFACE, line=LINE)
        add_rect(slide, 7.2, top, 0.07, 0.83, accent, rounded=False)
        add_text(slide, title, 7.47, top + 0.12, 3.0, 0.25, size=12, color=NAVY, bold=True)
        add_text(slide, label, 10.42, top + 0.12, 1.65, 0.22, size=9, color=accent, bold=True, align=PP_ALIGN.RIGHT)
        add_text(slide, detail, 7.47, top + 0.46, 4.5, 0.23, size=9.5, color=GRAY)
    footer(slide, 2)

    # 3. Editable feature matrix
    slide = prs.slides.add_slide(blank); background(slide, WHITE); header(slide, "Capability-by-capability evidence", "Competitive tracker", 3)
    table_shape = slide.shapes.add_table(7, 6, Inches(0.62), Inches(1.48), Inches(12.1), Inches(5.32))
    table = table_shape.table
    widths = [2.45, 1.95, 1.45, 1.1, 0.75, 4.4]
    for column, width in zip(table.columns, widths, strict=True): column.width = Inches(width)
    headers = ["FEATURE", "STRATACORP", "US", "GAP", "PRIORITY", "NOTES"]
    for column, value in enumerate(headers):
        cell = table.cell(0, column); cell.text = value; cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
        for run in cell.text_frame.paragraphs[0].runs: run.font.name = "Aptos"; run.font.size = Pt(9); run.font.bold = True; run.font.color.rgb = WHITE
    status_color = {"Parity": MUTED, "Behind": RED, "Ahead": GREEN, "Gap": AMBER}
    for row_index, row in enumerate(rows, 1):
        values = [row["feature_name"], row["stratacorp_has"], row["we_have"], row["gap_status"], row["priority"], row["notes"] or "—"]
        for column, value in enumerate(values):
            cell = table.cell(row_index, column); cell.text = str(value); cell.fill.solid(); cell.fill.fore_color.rgb = WHITE if row_index % 2 else SURFACE
            cell.margin_left = cell.margin_right = Inches(0.06)
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Aptos"; run.font.size = Pt(8.8); run.font.color.rgb = status_color[row["gap_status"]] if column == 3 else GRAY
                    if column in {0, 3, 4}: run.font.bold = True
        table.rows[row_index].height = Inches(0.75)
    table.rows[0].height = Inches(0.55)
    footer(slide, 3)

    # 4. External signal mosaic
    slide = prs.slides.add_slide(blank); background(slide, SURFACE); header(slide, "What the market is telling us", "Analyst + field intelligence", 4)
    add_rect(slide, 0.72, 1.55, 5.72, 4.98, WHITE, line=LINE)
    pill(slide, "ANALYST NOTE", 1.02, 1.84, 1.55, PURPLE)
    metrics = [("3x", "faster report generation"), ("4.2/5", "mobile beta rating"), ("15-20%", "mid-market undercut"), ("12 vs 30+", "workflow trigger types")]
    for index, (value, label) in enumerate(metrics):
        row, column = divmod(index, 2); left = 1.02 + column * 2.55; top = 2.4 + row * 1.27
        add_text(slide, value, left, top, 2.25, 0.42, size=22, color=NAVY, bold=True)
        add_text(slide, label, left, top + 0.5, 2.25, 0.32, size=9.5, color=MUTED)
    add_text(slide, "Pricing: Starter $29/user/mo  ·  Pro $79/user/mo  ·  Enterprise custom pricing", 1.02, 5.13, 4.95, 0.62, size=11, color=GRAY, bold=True)
    add_text(slide, "Embedded Analytics is now GA via DataViz Co.", 1.02, 5.88, 4.95, 0.28, size=10.5, color=PURPLE, bold=True)
    add_rect(slide, 6.86, 1.55, 5.75, 4.98, WHITE, line=LINE)
    pill(slide, "FIELD / SLACK", 7.16, 1.84, 1.55, BLUE)
    field_cards = [
        ("Meridian Corp", "30% Pro discount to switch  ·  ~$45K ARR", "Jordan Mills"),
        ("Mobile urgency", "Two prospects raised the app; one called it a deciding factor", "Casey Park"),
        ("NovaTech win-back", "Likes AI insights; custom workflows are 'basically nonexistent'", "Dana Reeves"),
    ]
    for index, (title, detail, owner) in enumerate(field_cards):
        top = 2.4 + index * 1.18
        add_text(slide, title, 7.18, top, 4.65, 0.3, size=13, color=NAVY, bold=True)
        add_text(slide, detail, 7.18, top + 0.38, 4.65, 0.45, size=10.5, color=GRAY)
        add_text(slide, owner, 11.25, top, 0.95, 0.22, size=8.5, color=MUTED, align=PP_ALIGN.RIGHT)
    add_text(slide, "NovaTech may return if we ship our AI module.", 7.18, 5.97, 4.7, 0.28, size=10.5, color=GREEN, bold=True)
    footer(slide, 4)

    # 5. Response priorities
    slide = prs.slides.add_slide(blank); background(slide, WHITE); header(slide, "Recommended response", "Action architecture", 5)
    priorities = [
        ("P0", "AI-Powered Insights", "Prepare battle cards; accelerate credible AI module proof points.", RED),
        ("P1", "Native Mobile App", "Give field teams a clear roadmap story for the Q3 2026 mobile plan.", AMBER),
        ("P1", "Embedded Analytics", "Equip RFP teams to answer the DataViz Co. / GA differentiator.", AMBER),
        ("DEFEND", "Custom Workflows", "Lead with our workflow advantage and customer success stories.", GREEN),
    ]
    for index, (tag, title, detail, accent) in enumerate(priorities):
        top = 1.55 + index * 1.12
        add_rect(slide, 0.72, top, 7.45, 0.92, SURFACE, line=LINE)
        pill(slide, tag, 0.98, top + 0.26, 0.82, accent)
        add_text(slide, title, 2.02, top + 0.16, 2.35, 0.3, size=14, color=NAVY, bold=True)
        add_text(slide, detail, 4.33, top + 0.16, 3.45, 0.5, size=10, color=GRAY)
    add_rect(slide, 8.58, 1.55, 3.9, 4.28, NAVY, line=NAVY)
    add_text(slide, "Pipeline watchlist", 8.93, 1.92, 3.15, 0.35, size=18, color=WHITE, bold=True)
    add_text(slide, "Meridian Corp", 8.93, 2.55, 2.7, 0.28, size=13, color=CYAN, bold=True)
    add_text(slide, "Defend ~$45K ARR opportunity against a 30% Pro discount.", 8.93, 2.94, 2.95, 0.62, size=10.5, color=RGBColor(0xD5, 0xDE, 0xE8))
    add_text(slide, "NovaTech", 8.93, 3.82, 2.7, 0.28, size=13, color=CYAN, bold=True)
    add_text(slide, "Win-back signal depends on AI module credibility; emphasize workflow advantage.", 8.93, 4.2, 2.95, 0.68, size=10.5, color=RGBColor(0xD5, 0xDE, 0xE8))
    add_text(slide, "Owner: Strategy + Product Marketing", 8.93, 5.25, 2.95, 0.25, size=9, color=RGBColor(0xA9, 0xB8, 0xC9), bold=True)
    footer(slide, 5)

    # 6. Pricing governance and delivery
    slide = prs.slides.add_slide(blank); background(slide, NAVY)
    add_text(slide, "PRICING GOVERNANCE  /  CONTROLLED EXECUTION", 0.72, 0.58, 7.5, 0.25, size=10, color=CYAN, bold=True)
    add_text(slide, "Capture the pricing signal.\nDo not bypass approval.", 0.72, 1.02, 7.1, 1.15, size=31, color=WHITE, bold=True)
    tiers = [("Starter", "$29/user/mo"), ("Pro", "$79/user/mo"), ("Enterprise", "custom pricing")]
    for index, (tier, price) in enumerate(tiers):
        left = 0.72 + index * 2.42
        add_rect(slide, left, 2.58, 2.12, 1.2, RGBColor(0x13, 0x36, 0x52), line=RGBColor(0x2A, 0x53, 0x72))
        add_text(slide, tier, left + 0.18, 2.83, 1.75, 0.24, size=10, color=RGBColor(0xA9, 0xC7, 0xDC), bold=True)
        add_text(slide, price, left + 0.18, 3.22, 1.75, 0.3, size=15, color=WHITE, bold=True)
    add_rect(slide, 0.72, 4.18, 7.0, 1.42, RGBColor(0x2C, 0x20, 0x2C), line=RGBColor(0x6F, 0x38, 0x3E))
    add_text(slide, "DO NOT directly change the pricing comparison page", 1.02, 4.52, 6.3, 0.3, size=15, color=RGBColor(0xFF, 0xB4, 0xAB), bold=True)
    add_text(slide, "Pricing Committee approval is required; reviews occur every Wednesday.", 1.02, 5.0, 6.25, 0.32, size=11, color=RGBColor(0xE8, 0xD7, 0xD9))
    add_rect(slide, 8.38, 1.02, 4.15, 4.58, WHITE, line=WHITE)
    add_text(slide, "NEXT CONTROLLED ACTION", 8.73, 1.38, 3.35, 0.24, size=10, color=BLUE, bold=True)
    add_text(slide, "Open HubSpot ticket", 8.73, 1.88, 3.3, 0.35, size=19, color=NAVY, bold=True)
    add_text(slide, "Subject", 8.73, 2.55, 1.1, 0.22, size=9, color=MUTED, bold=True)
    add_text(slide, "Pricing Page Update Request", 8.73, 2.9, 3.2, 0.54, size=13, color=INK, bold=True)
    add_text(slide, "Description must include", 8.73, 3.67, 1.9, 0.22, size=9, color=MUTED, bold=True)
    add_text(slide, "StrataCorp tiers, exact prices, and comparison-page changes requested", 8.73, 4.02, 3.15, 0.72, size=11, color=GRAY)
    add_text(slide, "Deliver brief to", 8.73, 4.95, 1.3, 0.22, size=9, color=MUTED, bold=True)
    add_text(slide, context["deliver_to"], 8.73, 5.28, 3.15, 0.28, size=11, color=BLUE, bold=True)
    footer(slide, 6, dark=True)

    output.parent.mkdir(parents=True, exist_ok=True); prs.save(output)
    if len(Presentation(output).slides) != 6: raise ValueError("Golden 保存后不是 6 页")


def main():
    parser = argparse.ArgumentParser(); parser.add_argument("--input-dir", type=Path, required=True); parser.add_argument("--output", type=Path, required=True); args = parser.parse_args()
    build(args.input_dir, args.output); print(f"{args.output}\nsha256={sha256(args.output)}")


if __name__ == "__main__": main()
