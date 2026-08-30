#!/usr/bin/env python3
"""AutomationBench 1115 多应用状态的逐源静态化。"""

from __future__ import annotations

import argparse
import json
from email.message import EmailMessage
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


TRACKER_ROWS = [
    [2, "Real-time Collaboration", "Yes", "Yes", "Parity", "P2", ""],
    [3, "AI-Powered Insights", "Yes (new in v4)", "Beta", "Behind", "P0", "Their implementation uses GPT-4 integration natively"],
    [4, "Custom Workflows", "Limited", "Yes", "Ahead", "P1", ""],
    [5, "SSO / Enterprise Auth", "Yes", "Yes", "Parity", "P3", ""],
    [6, "Native Mobile App", "Yes (new in v4)", "No", "Gap", "P1", "Our mobile roadmap is Q3 2026"],
    [7, "Embedded Analytics", "Yes", "Partial", "Behind", "P1", "They acquired DataViz Co. for this"],
]

ANALYST_BODY = """ANALYST BRIEF - EMBARGOED UNTIL JAN 27 9AM ET

StrataCorp today announced v4 of their platform. Key takeaways:

1. AI-Powered Insights is the headline feature, leveraging a deep GPT-4 integration. Early demos show 3x faster report generation vs their v3.

2. New mobile app addresses a long-standing gap. Initial reviews from beta testers are positive (4.2/5 rating).

3. Pricing restructured: Starter at $29/user/mo, Pro at $79/user/mo, Enterprise at custom pricing. This undercuts most mid-market competitors by 15-20%.

4. Embedded analytics powered by their DataViz Co. acquisition is now GA. Expect this to be a key differentiator in RFPs.

5. Weak spot: workflow customization remains limited. Their no-code builder supports only 12 trigger types vs. industry standard of 30+.

Recommendation: Monitor pipeline deals where StrataCorp is shortlisted. Prepare battle cards emphasizing workflow flexibility and existing customer success stories."""

POLICY_BODY = """Hi everyone,

Quick reminder: all pricing page changes require Pricing Committee approval. This includes competitor pricing comparisons, tier restructuring, and promotional offers.

If you need a pricing page update:
1. Open a HubSpot ticket with subject 'Pricing Page Update Request'
2. Include the specific changes needed in the description
3. The Pricing Committee reviews requests every Wednesday

Do NOT make direct changes to the pricing comparison page.

Thanks,
Product Lead"""

SLACK = {
    "channel": {"id": "C_cintel", "name": "competitive-intel", "is_private": False, "is_archived": False},
    "users": [
        {"id": "U_ae1", "name": "Jordan Mills", "email": "jordan@company.example.com"},
        {"id": "U_ae2", "name": "Casey Park", "email": "casey@company.example.com"},
        {"id": "U_ae3", "name": "Dana Reeves", "email": "dana@company.example.com"}
    ],
    "messages": [
        {"ts": "1738000000.000001", "user_id": "U_ae1", "text": "Just got off a call with Meridian Corp. They said StrataCorp offered them a 30% discount on the Pro tier to switch. Deal size was ~$45K ARR."},
        {"ts": "1738000000.000002", "user_id": "U_ae2", "text": "FYI their new mobile app is generating buzz. Two prospects mentioned it on calls today -- one said it was a deciding factor. We need a mobile story."},
        {"ts": "1738000000.000003", "user_id": "U_ae3", "text": "Talked to a churned customer (NovaTech) who moved to StrataCorp last quarter. They love the AI insights but complained that custom workflows are 'basically nonexistent.' They may come back if we ship our AI module."}
    ]
}

CONTEXT = {
    "schema_version": "0.1.0",
    "current_time": "2026-01-27T09:00:00Z",
    "deliver_to": "strategy@company.example.com",
    "required_action": {"type": "HubSpot ticket", "subject": "Pricing Page Update Request"},
    "prohibited_action": "Do NOT make direct changes to the pricing comparison page.",
    "adaptation": "Create a competitive-response PowerPoint and a pricing-governance action slide instead of writing to external systems."
}


def build_tracker(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Feature Comparison"
    headers = ["row_id", "feature_name", "stratacorp_has", "we_have", "gap_status", "priority", "notes"]
    sheet.append(headers)
    for row in TRACKER_ROWS:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(horizontal="center")
    widths = {"A": 9, "B": 30, "C": 22, "D": 16, "E": 15, "F": 10, "G": 48}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        row[6].alignment = Alignment(wrap_text=True)
    sheet.freeze_panes = "A2"
    table = Table(displayName="CompetitiveFeatureTracker", ref="A1:G7")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
    sheet.add_table(table)
    workbook.save(path)


def write_eml(path: Path, *, sender: str, recipient: str, subject: str, date: str, body: str) -> None:
    message = EmailMessage()
    message["From"] = sender
    message["To"] = recipient
    message["Subject"] = subject
    message["Date"] = date
    message.set_content(body)
    path.write_bytes(message.as_bytes())


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    build_tracker(args.output_dir / "Competitive_Intelligence_Tracker.xlsx")
    write_eml(args.output_dir / "Analyst_Note_StrataCorp_v4.eml", sender="research@analystfirm.example.com", recipient="me", subject="Quick Take: StrataCorp v4 Platform Launch", date="Tue, 27 Jan 2026 06:00:00 +0000", body=ANALYST_BODY)
    write_eml(args.output_dir / "Pricing_Page_Change_Policy.eml", sender="product-lead@company.example.com", recipient="marketing-all@company.example.com", subject="Reminder: Pricing Page Change Process", date="Tue, 20 Jan 2026 11:00:00 +0000", body=POLICY_BODY)
    (args.output_dir / "Competitive_Intel_Slack.json").write_text(json.dumps(SLACK, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (args.output_dir / "AutomationBench_1115_Context.json").write_text(json.dumps(CONTEXT, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
