#!/usr/bin/env python3
"""AutomationBench 1115 PPT Judge 校准样例。"""
from __future__ import annotations
import argparse, shutil, zipfile
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches
OUTPUT_NAME = "StrataCorp_v4_Competitive_Response_Brief.pptx"
C = "http://schemas.openxmlformats.org/drawingml/2006/chart"


def out(root, name):
    p = root / name / OUTPUT_NAME
    p.parent.mkdir(parents=True, exist_ok=True)
    return p
def mutate(g,o,fn): p=Presentation(g); fn(p); p.save(o)
def shape_on(prs,slide,token):
    m=[s for s in prs.slides[slide].shapes if getattr(s,"has_text_frame",False) and token in s.text]
    if not m: raise ValueError(token)
    return m[0]
def table(prs): return next(s.table for sl in prs.slides for s in sl.shapes if s.has_table and len(s.table.rows)==7)
def main():
    ap=argparse.ArgumentParser(); ap.add_argument("--golden",type=Path,required=True); ap.add_argument("--outdir",type=Path,required=True); a=ap.parse_args(); (a.outdir/"empty").mkdir(parents=True); out(a.outdir,"corrupt").write_bytes(b"bad\n")
    with zipfile.ZipFile(a.golden) as r,zipfile.ZipFile(out(a.outdir,"valid-repackage"),"w",compression=zipfile.ZIP_STORED) as w:
        for i in r.infolist(): w.writestr(i.filename,r.read(i.filename))
    def alt(p): p.slides[0].background.fill.solid(); p.slides[0].background.fill.fore_color.rgb=RGBColor(0x13,0x30,0x4A)
    mutate(a.golden,out(a.outdir,"valid-alt-cover"),alt)
    def chart_wrong(p):
        ch=next(s.chart for sl in p.slides for s in sl.shapes if s.shape_type==MSO_SHAPE_TYPE.CHART); cache=ch._chartSpace.find(f".//{{{C}}}numCache"); cache.findall(f"{{{C}}}pt")[0].find(f"{{{C}}}v").text="9"
    mutate(a.golden,out(a.outdir,"wrong-chart"),chart_wrong)
    def missing_row(p):
        t=table(p)
        for c in t.rows[2].cells: c.text=""
    mutate(a.golden,out(a.outdir,"missing-feature"),missing_row)
    def wrong_gap(p): table(p).cell(2,3).text="Parity"
    mutate(a.golden,out(a.outdir,"wrong-gap"),wrong_gap)
    def miss_analyst(p):
        s=shape_on(p,3,"3x"); s.text=s.text.replace("3x","2x")
    mutate(a.golden,out(a.outdir,"missing-analyst-fact"),miss_analyst)
    def miss_slack(p):
        s=shape_on(p,3,"~$45K ARR"); s.text=s.text.replace("~$45K ARR","~$50K ARR")
    mutate(a.golden,out(a.outdir,"wrong-slack-fact"),miss_slack)
    def wrong_priority(p): shape_on(p,4,"P0").text="P2"
    mutate(a.golden,out(a.outdir,"wrong-response-priority"),wrong_priority)
    def missing_defense(p): shape_on(p,4,"DEFEND").text=""
    mutate(a.golden,out(a.outdir,"missing-workflow-defense"),missing_defense)
    def wrong_price(p):
        s=shape_on(p,5,"$29/user/mo"); s.text="$39/user/mo"
    mutate(a.golden,out(a.outdir,"wrong-price"),wrong_price)
    def direct_change(p):
        s=shape_on(p,5,"DO NOT directly change"); s.text="Directly change the pricing comparison page"
    mutate(a.golden,out(a.outdir,"policy-violation"),direct_change)
    def wrong_ticket(p): shape_on(p,5,"Pricing Page Update Request").text="Update Pricing Now"
    mutate(a.golden,out(a.outdir,"wrong-ticket-subject"),wrong_ticket)
    def wrong_recipient(p): shape_on(p,5,"strategy@company.example.com").text="sales@company.example.com"
    mutate(a.golden,out(a.outdir,"wrong-recipient"),wrong_recipient)
    image=a.outdir/"flat.png"; Image.new("RGB",(1600,900),"white").save(image)
    def flat(p):
        sl=p.slides[2]
        for s in list(sl.shapes): s._element.getparent().remove(s._element)
        sl.shapes.add_picture(str(image),0,0,width=p.slide_width,height=p.slide_height)
    mutate(a.golden,out(a.outdir,"flattened-matrix"),flat)
    def hard(p): b=p.slides[-1].shapes.add_textbox(Inches(.1),Inches(.1),Inches(2),Inches(.2)); b.text="solution/Golden"
    mutate(a.golden,out(a.outdir,"hardcoded-marker"),hard)
    extra=a.outdir/"extra-output"; extra.mkdir(); shutil.copy2(a.golden,extra/OUTPUT_NAME); (extra/"debug.txt").write_text("debug\n")
if __name__=="__main__": main()
