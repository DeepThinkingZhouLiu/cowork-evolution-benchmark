#!/usr/bin/env python3
"""AutomationBench 1115 PPT 适配的任务专属确定性 Judge。"""
from __future__ import annotations
import argparse, hashlib, json, posixpath, re, subprocess, tempfile, zipfile
from pathlib import Path
from collections import Counter
from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

OUTPUT_NAME="StrataCorp_v4_Competitive_Response_Brief.pptx"
INPUT_HASHES={
"Competitive_Intelligence_Tracker.xlsx":"4bb463f0b76e21a875586f2f30ce666cde0a98280fa1d4ab281fa6d19c4f8883",
"Analyst_Note_StrataCorp_v4.eml":"58854e7344764bef72b7e4b0ae1fd9313b3ad3308293792292da2d50f445bdee",
"Pricing_Page_Change_Policy.eml":"36fd563f10d5a777d7c6ca90db270961011a3bab3481dea849ace4c69135dd7d",
"Competitive_Intel_Slack.json":"f9fc7f315c4c4ba7fddf18dd4a64f8225846a45fa146921cad50938624d8b665",
"AutomationBench_1115_Context.json":"6969d63a0d6163bbdaca24ac0495b73f3e0875997f54bf5143564fe3ab0102ed"}
TASK_ID="cowork-evo/automationbench-1115-stratacorp-competitive-response-deck"; RUBRIC_VERSION="automationbench-1115-adaptation-v0.1"; JUDGE_VERSION="automationbench-1115-deterministic-v0.1"
ANALYST_TOKENS=["ANALYST NOTE","3x","4.2/5","$29/user/mo","$79/user/mo","custom pricing","15-20%","12 vs 30+","Embedded Analytics","GA","DataViz Co."]
SLACK_TOKENS=["FIELD / SLACK","Meridian Corp","30%","~$45K ARR","Two prospects","deciding factor","NovaTech","basically nonexistent","AI module"]
PRICING_TOKENS=["$29/user/mo","$79/user/mo","custom pricing","DO NOT directly change the pricing comparison page","Pricing Committee","every Wednesday","Pricing Page Update Request","strategy@company.example.com","StrataCorp"]
SUSPICIOUS=["solution/golden","judge-result","artifact_dir","automationbench-1115","/root/"]
SPECS={
"R001":("PPTX 完整、16:9 且 6 页",6,"hurdle"),"R002":("原生图表的 gap status 为 2/2/1/1",10,"hurdle"),"R003":("6 行 feature 矩阵映射完整",22,"hurdle"),"R004":("分析师事实完整标源",14,"hurdle"),"R005":("Slack 事实完整标源",14,"hurdle"),"R006":("响应优先级与优势正确",10,"hurdle"),"R007":("价格治理、ticket 和受众正确",12,"hurdle"),"R008":("来源分隔且原生可编辑，无压平/标记/宏",10,"hurdle"),"R009":("6 页可渲染且输出干净",2,"positive"),"R010":("无 tracker、外部事实、优先级、治理或压平严重错误",-20,"penalty")}

def sha256(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def norm(v): return re.sub(r"\s+"," ",str(v).replace("\u00a0"," ").replace("\v"," ")).strip()
def locate(name):
    for p in (Path("/app/input_files")/name,Path(__file__).resolve().parents[1]/"data/input_files"/name):
        if p.is_file(): return p
    return None
def package_checks(path):
    e={"broken_member":None,"missing_targets":[],"macros":[]}
    try:
        with zipfile.ZipFile(path) as z:
            e["broken_member"]=z.testzip(); names=set(z.namelist()); e["macros"]=[n for n in names if n.lower().endswith("vbaproject.bin")]
            for rn in sorted(n for n in names if n.endswith(".rels")):
                root=etree.fromstring(z.read(rn)); owner="" if rn=="_rels/.rels" else posixpath.join(rn.split("/_rels/",1)[0],rn.split("/_rels/",1)[1][:-5]); base=posixpath.dirname(owner)
                for rel in root:
                    if rel.attrib.get("TargetMode")=="External": continue
                    t=rel.attrib.get("Target",""); resolved=t.lstrip("/") if t.startswith("/") else posixpath.normpath(posixpath.join(base,t))
                    if resolved not in names: e["missing_targets"].append(f"{rn} -> {resolved}")
        return e["broken_member"] is None and not e["missing_targets"] and not e["macros"],e
    except Exception as x: e["exception"]=str(x); return False,e
def expected_rows():
    for n,h in INPUT_HASHES.items():
        p=locate(n)
        if p is None or sha256(p)!=h: raise ValueError(f"冻结输入缺失或哈希错：{n}")
    wb=load_workbook(locate("Competitive_Intelligence_Tracker.xlsx"),data_only=True); ws=wb["Feature Comparison"]; headers=[c.value for c in ws[1]]; result={}
    for row in ws.iter_rows(min_row=2):
        d=dict(zip(headers,[c.value for c in row],strict=True)); result[norm(d["feature_name"])]=(norm(d["stratacorp_has"]),norm(d["we_have"]),norm(d["gap_status"]),norm(d["priority"]),norm(d["notes"]) if d["notes"] is not None else "—")
    return result
def deck_text(prs): return "\n".join(s.text for sl in prs.slides for s in sl.shapes if getattr(s,"has_text_frame",False))
def slide_text(sl): return "\n".join(s.text for s in sl.shapes if getattr(s,"has_text_frame",False) and not s.has_table)
def chart_check(prs):
    matches=[]; charts=[]
    for sl in prs.slides:
        for s in sl.shapes:
            if s.shape_type!=MSO_SHAPE_TYPE.CHART: continue
            c=s.chart; charts.append(c)
            try:
                cats=[norm(x.label) for x in c.plots[0].categories]; vals=[float(x) for x in c.series[0].values]; mapping=dict(zip(cats,vals,strict=True))
                if mapping=={"Parity":2.0,"Behind":2.0,"Ahead":1.0,"Gap":1.0}: matches.append(mapping)
            except Exception: pass
    return len(matches)==1,{"charts":len(charts),"matches":matches}
def table_check(prs,expected):
    candidates=[]
    for sn,sl in enumerate(prs.slides,1):
        for s in sl.shapes:
            if not s.has_table: continue
            t=s.table; hdr=[norm(c.text).upper() for c in t.rows[0].cells] if len(t.rows) else []; required=["FEATURE","STRATACORP","US","GAP","PRIORITY","NOTES"]
            if not all(x in hdr for x in required): continue
            pos={x:hdr.index(x) for x in required}; obs={}
            for ri in range(1,len(t.rows)):
                r=t.rows[ri]; feature=norm(r.cells[pos["FEATURE"]].text)
                if feature: obs[feature]=tuple(norm(r.cells[pos[x]].text) or "—" for x in required[1:])
            candidates.append((sn,obs))
    return sum(obs==expected for _,obs in candidates)==1,{"candidates":[{"slide":sn,"rows":obs} for sn,obs in candidates]}
def render(path,out):
    out.mkdir(parents=True,exist_ok=True); profile=out/"profile"; profile.mkdir(exist_ok=True); c=subprocess.run(["soffice",f"-env:UserInstallation=file://{profile}","--headless","--convert-to","pdf","--outdir",str(out),str(path)],stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,timeout=300); pdf=out/f"{path.stem}.pdf"
    if c.returncode or not pdf.is_file(): return 0,c.stdout[-500:]
    info=subprocess.run(["pdfinfo",str(pdf)],stdout=subprocess.PIPE,text=True).stdout; m=re.search(r"^Pages:\s+(\d+)",info,re.M); return int(m.group(1)) if m else 0,c.stdout[-500:]
def item(i,p,e):
    d,w,k=SPECS[i]; return {"criterion_id":i,"score":1.0 if p else 0.0,"evidence":e,"method":"deterministic","raw":{"description":d,"weight":w,"type":k,"passed":bool(p)}}
def result(path,criteria):
    hurdles=[x["criterion_id"] for x in criteria if x["raw"]["type"]=="hurdle" and not x["raw"]["passed"]]; total=sum(x["raw"]["weight"] for x in criteria if x["raw"]["weight"]>0); earned=sum(x["raw"]["weight"] for x in criteria if x["raw"]["weight"]>0 and x["raw"]["passed"]); penalty=sum(abs(x["raw"]["weight"]) for x in criteria if x["raw"]["weight"]<0 and not x["raw"]["passed"]); reward=0.0 if hurdles else max(0,min(1,(earned-penalty)/total)); return {"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":[sha256(path)] if path.is_file() else [],"criterion_results":criteria,"reward":reward,"passed":reward>=.8 and not hurdles,"hurdle_failures":hurdles}
def evaluate(path,artifact):
    try: expected=expected_rows(); prs=Presentation(path)
    except Exception as x: return result(path,[item(i,False,f"打开失败：{x}") for i in SPECS])
    criteria=[]; pkg,pe=package_checks(path); ratio=prs.slide_width/prs.slide_height if prs.slide_height else 0; f_ok=pkg and len(prs.slides)==6 and abs(ratio-16/9)<.02; criteria.append(item("R001",f_ok,f"页={len(prs.slides)} 比例={ratio:.4f} 包={json.dumps(pe,ensure_ascii=False)}"))
    c_ok,cd=chart_check(prs); criteria.append(item("R002",c_ok,json.dumps(cd,ensure_ascii=False)))
    t_ok,td=table_check(prs,expected); criteria.append(item("R003",t_ok,json.dumps(td,ensure_ascii=False)))
    analyst_slides=[i+1 for i,sl in enumerate(prs.slides) if all(tok in slide_text(sl) for tok in ANALYST_TOKENS)]; a_ok=bool(analyst_slides); criteria.append(item("R004",a_ok,f"分析师事实页={analyst_slides}"))
    slack_slides=[i+1 for i,sl in enumerate(prs.slides) if all(tok in slide_text(sl) for tok in SLACK_TOKENS)]; s_ok=bool(slack_slides); criteria.append(item("R005",s_ok,f"Slack 事实页={slack_slides}"))
    response=[slide_text(sl) for sl in prs.slides if "Recommended response" in slide_text(sl)]; needed=["AI-Powered Insights","P0","Native Mobile App","Embedded Analytics","P1","Custom Workflows","DEFEND","battle cards","advantage"]; r_ok=len(response)==1 and all(tok in response[0] for tok in needed); criteria.append(item("R006",r_ok,f"响应页={len(response)} 缺失={[x for x in needed if not response or x not in response[0]]}"))
    pricing=[slide_text(sl) for sl in prs.slides if "PRICING GOVERNANCE" in slide_text(sl)]; p_ok=len(pricing)==1 and all(tok in pricing[0] for tok in PRICING_TOKENS); criteria.append(item("R007",p_ok,f"价格页={len(pricing)} 缺失={[x for x in PRICING_TOKENS if not pricing or x not in pricing[0]]}"))
    text=deck_text(prs); tables=sum(s.has_table for sl in prs.slides for s in sl.shapes); charts=sum(s.shape_type==MSO_SHAPE_TYPE.CHART for sl in prs.slides for s in sl.shapes); full=[s.shape_id for sl in prs.slides for s in sl.shapes if s.shape_type==MSO_SHAPE_TYPE.PICTURE and s.width*s.height>=prs.slide_width*prs.slide_height*.8]; media=[]
    with zipfile.ZipFile(path) as z: media=[n for n in z.namelist() if n.startswith("ppt/media/") and not n.endswith("/")]
    markers=[m for m in SUSPICIOUS if m in text.lower()]; source_tags="ANALYST NOTE" in text and "FIELD / SLACK" in text; e_ok=tables>=1 and charts>=1 and not full and not media and not markers and source_tags and not pe.get("macros"); criteria.append(item("R008",e_ok,f"表={tables} 图={charts} 整页图={full} 媒体={media} 标记={markers} 源标签={source_tags}"))
    pages=0; log=""
    try:
        with tempfile.TemporaryDirectory(prefix="ab1115-render-") as d: pages,log=render(path,Path(d))
    except Exception as x: log=str(x)
    delivered=sorted(str(x.relative_to(artifact)) for x in artifact.rglob("*") if x.is_file()); h_ok=pages==6 and delivered==[OUTPUT_NAME] and pkg; criteria.append(item("R009",h_ok,f"页={pages} 文件={delivered} 日志={log}"))
    severe=not all([c_ok,t_ok,a_ok,s_ok,r_ok,p_ok,e_ok]); criteria.append(item("R010",not severe,f"严重错误={severe}")); return result(path,criteria)
def main():
    p=argparse.ArgumentParser(); p.add_argument("--artifact-dir",type=Path,required=True); p.add_argument("--result",type=Path,required=True); p.add_argument("--reward-file",type=Path,required=True); a=p.parse_args(); path=a.artifact_dir/OUTPUT_NAME; out=evaluate(path,a.artifact_dir) if path.is_file() else result(path,[item(i,False,f"缺失 {OUTPUT_NAME}") for i in SPECS]); a.result.parent.mkdir(parents=True,exist_ok=True); a.reward_file.parent.mkdir(parents=True,exist_ok=True); a.result.write_text(json.dumps(out,ensure_ascii=False,indent=2)+"\n"); a.reward_file.write_text(f"{out['reward']:.6f}\n"); print(json.dumps(out,ensure_ascii=False,indent=2))
if __name__=="__main__": main()
