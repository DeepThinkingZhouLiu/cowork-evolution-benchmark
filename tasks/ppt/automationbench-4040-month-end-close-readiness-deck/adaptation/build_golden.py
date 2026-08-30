#!/usr/bin/env python3
"""AutomationBench 4040 逐题设计的财务结账就绪简报 Golden。"""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


XLSX_SHA256 = "067ffd3a28b7608c0e0f7f00e34aa62c3c8a928981460c39e4e7e8cab1694b70"
CONTEXT_SHA256 = "76003f8de514e771e61103bb80f2e681d199ffd9bfda0980c076e7992ec5667b"

NAVY = RGBColor(0x0F, 0x27, 0x47)
BLUE = RGBColor(0x2F, 0x80, 0xED)
GREEN = RGBColor(0x12, 0xB7, 0x6A)
AMBER = RGBColor(0xF7, 0x90, 0x09)
RED = RGBColor(0xD9, 0x2D, 0x20)
GRAY_900 = RGBColor(0x1D, 0x29, 0x39)
GRAY_700 = RGBColor(0x34, 0x40, 0x54)
GRAY_500 = RGBColor(0x66, 0x70, 0x85)
GRAY_200 = RGBColor(0xEA, 0xEC, 0xF0)
GRAY_100 = RGBColor(0xF2, 0xF4, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def set_background(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(slide, text: str, left: float, top: float, width: float, height: float, *, size: float = 18, color: RGBColor = GRAY_900, bold: bool = False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font: str = "Aptos", margin: float = 0.04):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    return shape


def add_rect(slide, left: float, top: float, width: float, height: float, fill: RGBColor, *, radius: bool = True, line: RGBColor | None = None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def add_header(slide, title: str, kicker: str, page: int) -> None:
    add_text(slide, kicker.upper(), 0.65, 0.32, 7.4, 0.25, size=10, color=BLUE, bold=True)
    add_text(slide, title, 0.65, 0.62, 11.8, 0.56, size=27, color=NAVY, bold=True)
    add_text(slide, f"{page:02d} / 05", 11.7, 0.4, 0.9, 0.25, size=9, color=GRAY_500, align=PP_ALIGN.RIGHT)
    add_rect(slide, 0.65, 1.24, 12.03, 0.02, GRAY_200, radius=False)


def add_footer(slide, page: int, dark: bool = False) -> None:
    color = RGBColor(0x98, 0xA2, 0xB3) if not dark else RGBColor(0xB8, 0xC7, 0xD9)
    add_text(slide, "AutomationBench 4040  ·  Snapshot 2026-02-03 16:00Z", 0.65, 7.13, 8.0, 0.18, size=8, color=color)
    add_text(slide, f"{page}", 12.2, 7.13, 0.45, 0.18, size=8, color=color, align=PP_ALIGN.RIGHT)


def add_metric(slide, value: str, label: str, left: float, top: float, width: float, *, accent: RGBColor, dark: bool = False) -> None:
    fill = RGBColor(0x16, 0x36, 0x5F) if dark else WHITE
    line = RGBColor(0x2C, 0x4C, 0x70) if dark else GRAY_200
    add_rect(slide, left, top, width, 1.18, fill, line=line)
    add_rect(slide, left, top, 0.07, 1.18, accent, radius=False)
    add_text(slide, value, left + 0.24, top + 0.17, width - 0.35, 0.46, size=25, color=WHITE if dark else NAVY, bold=True)
    add_text(slide, label, left + 0.24, top + 0.72, width - 0.35, 0.25, size=10, color=RGBColor(0xD0, 0xD5, 0xDD) if dark else GRAY_500, bold=True)


def load_inputs(xlsx: Path, context_path: Path) -> tuple[list[dict], dict]:
    if sha256(xlsx) != XLSX_SHA256 or sha256(context_path) != CONTEXT_SHA256:
        raise ValueError("冻结输入 SHA256 不匹配")
    workbook = load_workbook(xlsx, data_only=True)
    sheet = workbook["January 2026 Checklist"]
    headers = [cell.value for cell in sheet[1]]
    rows = [dict(zip(headers, [cell.value for cell in row], strict=True)) for row in sheet.iter_rows(min_row=2)]
    context = json.loads(context_path.read_text(encoding="utf-8"))
    if len(rows) != 7 or context["snapshot_time"] != "2026-02-03T16:00:00Z":
        raise ValueError("输入行数或冻结时间不匹配")
    return rows, context


def build(xlsx: Path, context_path: Path, output: Path) -> None:
    rows, context = load_inputs(xlsx, context_path)
    complete = [row for row in rows if row["Status"] == "Complete"]
    incomplete = [row for row in rows if row["Status"] != "Complete"]
    if len(complete) != 4 or len(incomplete) != 3:
        raise ValueError("已审计的 4/3 完成分布不匹配")
    snapshot_date = datetime.fromisoformat(context["snapshot_time"].replace("Z", "+00:00")).date()

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    # 1. Executive cover
    slide = presentation.slides.add_slide(blank)
    set_background(slide, NAVY)
    add_text(slide, "MONTH-END CLOSE  /  JANUARY 2026", 0.72, 0.65, 8.0, 0.28, size=11, color=RGBColor(0x9B, 0xC2, 0xF5), bold=True)
    add_text(slide, "Close Readiness\nBriefing", 0.72, 1.12, 7.3, 1.55, size=38, color=WHITE, bold=True)
    add_text(slide, "Controller & CFO review  ·  As of 03 Feb 2026, 16:00 UTC", 0.75, 2.95, 7.0, 0.36, size=13, color=RGBColor(0xD0, 0xD5, 0xDD))
    add_rect(slide, 9.02, 0.85, 3.58, 2.72, RGBColor(0x16, 0x36, 0x5F), line=RGBColor(0x2C, 0x4C, 0x70))
    add_text(slide, "57.1%", 9.35, 1.17, 2.9, 0.82, size=43, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "COMPLETE", 9.35, 2.05, 2.9, 0.3, size=11, color=RGBColor(0xB8, 0xC7, 0xD9), bold=True, align=PP_ALIGN.CENTER)
    add_rect(slide, 9.64, 2.62, 2.3, 0.5, RED)
    add_text(slide, "NOT READY", 9.64, 2.69, 2.3, 0.28, size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_metric(slide, "7", "CHECKLIST ITEMS", 0.72, 4.8, 3.75, accent=BLUE, dark=True)
    add_metric(slide, "4", "COMPLETE", 4.79, 4.8, 3.75, accent=GREEN, dark=True)
    add_metric(slide, "3", "OPEN / BLOCKED", 8.86, 4.8, 3.75, accent=RED, dark=True)
    add_footer(slide, 1, dark=True)

    # 2. Readiness overview
    slide = presentation.slides.add_slide(blank)
    set_background(slide, WHITE)
    add_header(slide, "Readiness at a glance", "Executive status", 2)
    chart_data = ChartData()
    chart_data.categories = ["Complete", "Incomplete"]
    chart_data.add_series("Checklist items", [4, 3])
    chart_shape = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(0.72), Inches(1.58), Inches(5.0), Inches(4.72), chart_data)
    chart = chart_shape.chart
    chart.has_title = False
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    plot = chart.plots[0]
    plot.hole_size = 62
    plot.has_data_labels = True
    plot.data_labels.position = XL_DATA_LABEL_POSITION.BEST_FIT
    plot.data_labels.show_category_name = True
    plot.data_labels.show_value = True
    plot.series[0].points[0].format.fill.solid()
    plot.series[0].points[0].format.fill.fore_color.rgb = GREEN
    plot.series[0].points[1].format.fill.solid()
    plot.series[0].points[1].format.fill.fore_color.rgb = RED
    add_text(slide, "4 / 7", 2.03, 3.05, 2.35, 0.55, size=30, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "items complete", 2.03, 3.62, 2.35, 0.3, size=11, color=GRAY_500, align=PP_ALIGN.CENTER)
    summary = [("Complete", "4", GREEN), ("In Progress", "1", AMBER), ("Not Started", "1", RED), ("Blocked", "1", RED)]
    for index, (label, value, accent) in enumerate(summary):
        row, column = divmod(index, 2)
        left = 6.15 + column * 3.12
        top = 1.76 + row * 1.52
        add_metric(slide, value, label.upper(), left, top, 2.78, accent=accent)
    add_rect(slide, 6.15, 4.95, 5.9, 1.14, GRAY_100, line=GRAY_200)
    add_text(slide, "Decision signal", 6.42, 5.15, 1.5, 0.24, size=10, color=GRAY_500, bold=True)
    add_text(slide, "NOT READY", 7.98, 5.1, 1.6, 0.32, size=16, color=RED, bold=True)
    add_text(slide, "2 overdue items  ·  1 due today and blocked", 6.42, 5.58, 5.25, 0.27, size=12, color=GRAY_700)
    add_footer(slide, 2)

    # 3. Full editable checklist
    slide = presentation.slides.add_slide(blank)
    set_background(slide, WHITE)
    add_header(slide, "Full close checklist", "Source-of-truth detail", 3)
    table_shape = slide.shapes.add_table(8, 5, Inches(0.72), Inches(1.55), Inches(11.9), Inches(4.98))
    table = table_shape.table
    widths = [4.0, 2.05, 1.8, 1.45, 2.6]
    for column, width in zip(table.columns, widths, strict=True):
        column.width = Inches(width)
    headers = ["ITEM", "OWNER", "STATUS", "DUE", "READINESS FLAG"]
    for column, value in enumerate(headers):
        cell = table.cell(0, column)
        cell.text = value
        cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
        for paragraph in cell.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.LEFT
            for run in paragraph.runs:
                run.font.name = "Aptos"; run.font.size = Pt(10); run.font.bold = True; run.font.color.rgb = WHITE
    for row_index, row in enumerate(rows, 1):
        due = str(row["Due"])
        flag = "On track"
        if row["Status"] != "Complete":
            flag = "Overdue" if datetime.fromisoformat(due).date() < snapshot_date else "Due today / blocked"
        values = [row["Item"], row["Owner"], row["Status"], due, flag]
        for column, value in enumerate(values):
            cell = table.cell(row_index, column)
            cell.text = str(value)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE if row_index % 2 else GRAY_100
            cell.margin_left = cell.margin_right = Inches(0.08)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.LEFT
                for run in paragraph.runs:
                    run.font.name = "Aptos"; run.font.size = Pt(10.5); run.font.color.rgb = GRAY_700
                    if column in {2, 4} and row["Status"] != "Complete":
                        run.font.bold = True; run.font.color.rgb = RED if row["Status"] in {"Not Started", "Blocked"} else AMBER
        table.rows[row_index].height = Inches(0.6)
    table.rows[0].height = Inches(0.55)
    add_footer(slide, 3)

    # 4. Incomplete ownership cards
    slide = presentation.slides.add_slide(blank)
    set_background(slide, GRAY_100)
    add_header(slide, "Open items requiring ownership", "Exception review", 4)
    accent_by_status = {"In Progress": AMBER, "Not Started": RED, "Blocked": RED}
    for index, row in enumerate(incomplete):
        left = 0.72 + index * 4.05
        add_rect(slide, left, 1.62, 3.7, 4.75, WHITE, line=GRAY_200)
        accent = accent_by_status[row["Status"]]
        add_rect(slide, left, 1.62, 3.7, 0.09, accent, radius=False)
        priority = "OVERDUE" if datetime.fromisoformat(str(row["Due"])).date() < snapshot_date else "DUE TODAY / BLOCKED"
        add_text(slide, priority, left + 0.24, 1.93, 3.1, 0.24, size=9, color=accent, bold=True)
        add_text(slide, str(row["Item"]), left + 0.24, 2.3, 3.15, 0.82, size=19, color=NAVY, bold=True)
        add_text(slide, "OWNER", left + 0.24, 3.26, 0.75, 0.2, size=9, color=GRAY_500, bold=True)
        add_text(slide, str(row["Owner"]), left + 1.05, 3.2, 2.25, 0.3, size=12, color=GRAY_900, bold=True)
        add_text(slide, str(row["Email"]), left + 0.24, 3.62, 3.15, 0.27, size=9.5, color=GRAY_500)
        add_text(slide, "STATUS", left + 0.24, 4.16, 0.75, 0.2, size=9, color=GRAY_500, bold=True)
        add_text(slide, str(row["Status"]), left + 1.05, 4.1, 2.25, 0.3, size=12, color=accent, bold=True)
        add_text(slide, "DUE", left + 0.24, 4.62, 0.75, 0.2, size=9, color=GRAY_500, bold=True)
        add_text(slide, str(row["Due"]), left + 1.05, 4.56, 2.25, 0.3, size=12, color=GRAY_900)
        if row.get("Notes"):
            add_rect(slide, left + 0.24, 5.15, 3.15, 0.83, RGBColor(0xFF, 0xF4, 0xED), line=RGBColor(0xFE, 0xCD, 0xCA))
            add_text(slide, str(row["Notes"]), left + 0.42, 5.34, 2.82, 0.45, size=10, color=RED, bold=True)
        else:
            add_text(slide, "Immediate owner follow-up required.", left + 0.24, 5.32, 3.1, 0.38, size=10.5, color=GRAY_700)
    add_footer(slide, 4)

    # 5. Communication-ready conclusion
    slide = presentation.slides.add_slide(blank)
    set_background(slide, NAVY)
    add_text(slide, "COMMUNICATION-READY CLOSE UPDATE", 0.72, 0.58, 7.6, 0.25, size=10, color=RGBColor(0x9B, 0xC2, 0xF5), bold=True)
    add_text(slide, "January close is not ready", 0.72, 1.02, 7.7, 0.68, size=32, color=WHITE, bold=True)
    add_text(slide, "4 of 7 checklist items are complete (57.1%). Revenue accruals and intercompany elimination are overdue; tax provision estimate is blocked pending state tax guidance.", 0.75, 1.95, 7.35, 1.15, size=16, color=RGBColor(0xE4, 0xE7, 0xEC))
    add_rect(slide, 0.75, 3.55, 7.35, 2.1, RGBColor(0x16, 0x36, 0x5F), line=RGBColor(0x2C, 0x4C, 0x70))
    add_text(slide, "Owners to follow up", 1.02, 3.87, 2.2, 0.28, size=11, color=RGBColor(0x9B, 0xC2, 0xF5), bold=True)
    add_text(slide, "Carol Diaz  ·  Revenue accruals\nDave Kim  ·  Intercompany elimination\nGrace Wu  ·  Tax provision estimate", 1.02, 4.28, 6.3, 1.05, size=14, color=WHITE, bold=True)
    add_rect(slide, 8.72, 1.02, 3.82, 4.63, WHITE, line=WHITE)
    add_text(slide, "DELIVERY AUDIENCE", 9.02, 1.36, 3.0, 0.24, size=10, color=BLUE, bold=True)
    add_text(slide, "Email", 9.02, 1.85, 1.0, 0.25, size=10, color=GRAY_500, bold=True)
    add_text(slide, "controller@company.example.com\ncfo@company.example.com", 9.02, 2.2, 3.0, 0.72, size=12, color=GRAY_900, bold=True)
    add_text(slide, "Slack", 9.02, 3.22, 1.0, 0.25, size=10, color=GRAY_500, bold=True)
    add_text(slide, "#finance-alerts", 9.02, 3.58, 3.0, 0.32, size=14, color=NAVY, bold=True)
    add_text(slide, "Overall status", 9.02, 4.33, 1.4, 0.25, size=10, color=GRAY_500, bold=True)
    add_rect(slide, 9.02, 4.72, 2.25, 0.48, RED)
    add_text(slide, "NOT READY", 9.02, 4.79, 2.25, 0.27, size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 5, dark=True)

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    reopened = Presentation(output)
    if len(reopened.slides) != 5:
        raise ValueError("Golden 保存后不是 5 页")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--context", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.xlsx, args.context, args.output)
    print(f"{args.output}\nsha256={sha256(args.output)}")


if __name__ == "__main__":
    main()
