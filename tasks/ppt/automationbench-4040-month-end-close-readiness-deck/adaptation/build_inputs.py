#!/usr/bin/env python3
"""将 AutomationBench 4040 的已审计应用状态静态化为两个可交付输入。"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROWS = [
    [2, "Bank reconciliation", "Alice Park", "alice@company.example.com", "Complete", "2026-02-01", ""],
    [3, "AP cutoff review", "Bob Chen", "bob@company.example.com", "Complete", "2026-02-01", ""],
    [4, "Revenue accruals", "Carol Diaz", "carol@company.example.com", "In Progress", "2026-02-02", ""],
    [5, "Intercompany elimination", "Dave Kim", "dave@company.example.com", "Not Started", "2026-02-02", ""],
    [6, "Fixed asset depreciation", "Eve Liu", "eve@company.example.com", "Complete", "2026-02-02", ""],
    [7, "Payroll reconciliation", "Frank Torres", "frank@company.example.com", "Complete", "2026-02-03", ""],
    [8, "Tax provision estimate", "Grace Wu", "grace@company.example.com", "Blocked", "2026-02-03", "Waiting on state tax guidance"],
]

CONTEXT = {
    "schema_version": "0.1.0",
    "snapshot_time": "2026-02-03T16:00:00Z",
    "source_systems": {
        "google_sheets": {"spreadsheet_id": "ss_close", "title": "Month-End Close", "worksheet_id": "ws_jan_checklist", "worksheet_title": "January 2026 Checklist"},
        "slack": {"channel_id": "C_FIN_ALERTS", "channel_name": "finance-alerts", "is_private": False},
    },
    "report_audience": ["controller@company.example.com", "cfo@company.example.com"],
    "adaptation": "Create a PowerPoint close-readiness briefing instead of sending external messages.",
}


def build_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "January 2026 Checklist"
    headers = ["Row ID", "Item", "Owner", "Email", "Status", "Due", "Notes"]
    sheet.append(headers)
    for row in ROWS:
        sheet.append(row)
    header_fill = PatternFill("solid", fgColor="17365D")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for row in sheet.iter_rows(min_row=2):
        row[5].number_format = "yyyy-mm-dd"
        row[6].alignment = Alignment(wrap_text=True)
    widths = {"A": 10, "B": 30, "C": 18, "D": 32, "E": 16, "F": 14, "G": 36}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    table = Table(displayName="JanuaryCloseChecklist", ref="A1:G8")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
    sheet.add_table(table)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    build_xlsx(args.output_dir / "January_2026_Close_Checklist.xlsx")
    (args.output_dir / "AutomationBench_4040_Context.json").write_text(json.dumps(CONTEXT, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
