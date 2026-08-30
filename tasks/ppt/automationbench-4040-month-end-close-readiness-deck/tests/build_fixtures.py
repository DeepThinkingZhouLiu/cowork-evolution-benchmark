#!/usr/bin/env python3
"""AutomationBench 4040 PPT Judge 校准样例。"""

from __future__ import annotations

import argparse
import shutil
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches


OUTPUT_NAME = "January_2026_Close_Readiness_Briefing.pptx"
CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"


def output_path(root: Path, case: str) -> Path:
    path = root / case / OUTPUT_NAME
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def save_mutation(golden: Path, output: Path, callback) -> None:
    presentation = Presentation(golden)
    callback(presentation)
    presentation.save(output)


def repackage(golden: Path, output: Path) -> None:
    with zipfile.ZipFile(golden) as reader, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_STORED) as writer:
        for info in reader.infolist():
            writer.writestr(info.filename, reader.read(info.filename))


def text_shape(prs, token: str):
    matches = [shape for slide in prs.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False) and token in shape.text]
    if not matches:
        raise ValueError(f"缺少文字：{token}")
    return matches[0]


def detail_table(prs):
    matches = [shape.table for slide in prs.slides for shape in slide.shapes if shape.has_table and len(shape.table.rows) == 8]
    if len(matches) != 1:
        raise ValueError("缺少唯一 8 行明细表")
    return matches[0]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()
    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    repackage(args.golden, output_path(args.outdir, "valid-repackage"))

    def alt_theme(prs):
        prs.slides[0].background.fill.solid()
        prs.slides[0].background.fill.fore_color.rgb = RGBColor(0x18, 0x34, 0x5A)
    save_mutation(args.golden, output_path(args.outdir, "valid-alt-cover"), alt_theme)

    def wrong_summary(prs):
        for slide in prs.slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and "57.1%" in shape.text:
                    shape.text = shape.text.replace("57.1%", "71.4%")
    save_mutation(args.golden, output_path(args.outdir, "wrong-summary"), wrong_summary)

    def wrong_chart(prs):
        chart = next(shape.chart for slide in prs.slides for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.CHART)
        cache = chart._chartSpace.find(f".//{{{CHART_NS}}}numCache")
        points = cache.findall(f"{{{CHART_NS}}}pt")
        points[0].find(f"{{{CHART_NS}}}v").text = "5"
    save_mutation(args.golden, output_path(args.outdir, "wrong-chart-value"), wrong_chart)

    def chart_removed(prs):
        shape = next(shape for slide in prs.slides for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.CHART)
        shape._element.getparent().remove(shape._element)
    save_mutation(args.golden, output_path(args.outdir, "chart-removed"), chart_removed)

    def missing_item(prs):
        table = detail_table(prs)
        for cell in table.rows[3].cells:
            cell.text = ""
    save_mutation(args.golden, output_path(args.outdir, "missing-table-item"), missing_item)

    def wrong_owner(prs):
        table = detail_table(prs)
        table.cell(3, 1).text = "Alice Park"
    save_mutation(args.golden, output_path(args.outdir, "wrong-table-owner"), wrong_owner)

    def missing_note(prs):
        text_shape(prs, "Waiting on state tax guidance").text = ""
    save_mutation(args.golden, output_path(args.outdir, "missing-tax-note"), missing_note)

    def wrong_due_label(prs):
        text_shape(prs, "DUE TODAY / BLOCKED").text = "ON TRACK"
    save_mutation(args.golden, output_path(args.outdir, "wrong-due-classification"), wrong_due_label)

    def wrong_recipient(prs):
        shape = text_shape(prs, "controller@company.example.com")
        shape.text = shape.text.replace("controller@company.example.com", "ops@company.example.com")
    save_mutation(args.golden, output_path(args.outdir, "wrong-recipient"), wrong_recipient)

    def wrong_channel(prs):
        text_shape(prs, "#finance-alerts").text = "#general"
    save_mutation(args.golden, output_path(args.outdir, "wrong-channel"), wrong_channel)

    def invented_amount(prs):
        box = prs.slides[-1].shapes.add_textbox(Inches(0.8), Inches(6.2), Inches(3), Inches(0.3))
        box.text = "Estimated exposure: $1,000"
    save_mutation(args.golden, output_path(args.outdir, "invented-amount"), invented_amount)

    placeholder = args.outdir / "flatten-placeholder.png"
    Image.new("RGB", (1600, 900), "white").save(placeholder)
    def flattened(prs):
        slide = prs.slides[2]
        for shape in list(slide.shapes):
            shape._element.getparent().remove(shape._element)
        slide.shapes.add_picture(str(placeholder), 0, 0, width=prs.slide_width, height=prs.slide_height)
    save_mutation(args.golden, output_path(args.outdir, "flattened-detail-slide"), flattened)

    def hardcoded(prs):
        box = prs.slides[-1].shapes.add_textbox(Inches(0.1), Inches(0.1), Inches(2), Inches(0.2))
        box.text = "solution/Golden"
    save_mutation(args.golden, output_path(args.outdir, "hardcoded-marker"), hardcoded)

    extra = args.outdir / "extra-output"
    extra.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.golden, extra / OUTPUT_NAME)
    (extra / "debug.txt").write_text("debug\n", encoding="utf-8")


if __name__ == "__main__":
    main()
