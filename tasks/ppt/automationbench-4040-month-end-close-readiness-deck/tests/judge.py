#!/usr/bin/env python3
"""AutomationBench 4040 PPT 适配的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import posixpath
import re
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "January_2026_Close_Readiness_Briefing.pptx"
XLSX_NAME = "January_2026_Close_Checklist.xlsx"
CONTEXT_NAME = "AutomationBench_4040_Context.json"
XLSX_SHA256 = "067ffd3a28b7608c0e0f7f00e34aa62c3c8a928981460c39e4e7e8cab1694b70"
CONTEXT_SHA256 = "76003f8de514e771e61103bb80f2e681d199ffd9bfda0980c076e7992ec5667b"
TASK_ID = "cowork-evo/automationbench-4040-month-end-close-readiness-deck"
RUBRIC_VERSION = "automationbench-4040-adaptation-v0.1"
JUDGE_VERSION = "automationbench-4040-deterministic-v0.1"
INCOMPLETE = {
    "Revenue accruals": ("Carol Diaz", "In Progress", "2026-02-02", "OVERDUE"),
    "Intercompany elimination": ("Dave Kim", "Not Started", "2026-02-02", "OVERDUE"),
    "Tax provision estimate": ("Grace Wu", "Blocked", "2026-02-03", "DUE TODAY / BLOCKED"),
}
NOTE = "Waiting on state tax guidance"
RECIPIENTS = {"controller@company.example.com", "cfo@company.example.com"}
CHANNEL = "finance-alerts"
SUSPICIOUS_MARKERS = ["solution/golden", "judge-result", "artifact_dir", "automationbench-4040", "/root/"]

SPECS = {
    "R001": ("PPTX 完整、16:9 且恰好 5 页", 6, "hurdle"),
    "R002": ("Not Ready、7/4/3 和 57.1% 总体统计正确", 14, "hurdle"),
    "R003": ("原生完成/未完成图表为 4/3", 12, "hurdle"),
    "R004": ("7 行原生明细表映射完整准确", 22, "hurdle"),
    "R005": ("3 个未完成项与 Tax 备注已单独标出", 16, "hurdle"),
    "R006": ("2 个 overdue 和 1 个 due today/blocked 关联正确", 8, "hurdle"),
    "R007": ("Controller/CFO/finance-alerts 受众正确且无编造金额", 8, "hurdle"),
    "R008": ("图表、表格和文本可编辑且无压平/标记/宏", 8, "hurdle"),
    "R009": ("5 页可渲染且输出干净", 6, "positive"),
    "R010": ("无统计、明细、责任、日期、受众或压平严重错误", -20, "penalty"),
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", str(text).replace("\u00a0", " ").replace("\v", " ")).strip()


def locate_input(name: str) -> Path | None:
    candidates = [Path("/app/input_files") / name, Path(__file__).resolve().parents[1] / "data/input_files" / name]
    return next((path for path in candidates if path.is_file()), None)


def package_checks(path: Path) -> tuple[bool, dict]:
    evidence = {"broken_member": None, "missing_targets": [], "macros": []}
    try:
        with zipfile.ZipFile(path) as archive:
            evidence["broken_member"] = archive.testzip()
            names = set(archive.namelist())
            evidence["macros"] = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
            for rels_name in sorted(name for name in names if name.endswith(".rels")):
                root = etree.fromstring(archive.read(rels_name))
                if rels_name == "_rels/.rels":
                    owner = ""
                else:
                    prefix, filename = rels_name.split("/_rels/", 1)
                    owner = posixpath.join(prefix, filename[:-5])
                base = posixpath.dirname(owner)
                for relation in root:
                    if relation.attrib.get("TargetMode") == "External":
                        continue
                    target = relation.attrib.get("Target", "")
                    resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                    if resolved not in names:
                        evidence["missing_targets"].append(f"{rels_name} -> {resolved}")
        return evidence["broken_member"] is None and not evidence["missing_targets"] and not evidence["macros"], evidence
    except Exception as exc:
        evidence["exception"] = str(exc)
        return False, evidence


def load_expected() -> tuple[dict, dict]:
    xlsx = locate_input(XLSX_NAME)
    context_path = locate_input(CONTEXT_NAME)
    if xlsx is None or context_path is None or sha256(xlsx) != XLSX_SHA256 or sha256(context_path) != CONTEXT_SHA256:
        raise ValueError("冻结输入缺失或 SHA256 不匹配")
    workbook = load_workbook(xlsx, data_only=True)
    sheet = workbook["January 2026 Checklist"]
    headers = [cell.value for cell in sheet[1]]
    rows = [dict(zip(headers, [cell.value for cell in row], strict=True)) for row in sheet.iter_rows(min_row=2)]
    facts = {normalize(row["Item"]): (normalize(row["Owner"]), normalize(row["Status"]), normalize(row["Due"])) for row in rows}
    return facts, json.loads(context_path.read_text(encoding="utf-8"))


def all_text(presentation: Presentation) -> str:
    return "\n".join(shape.text for slide in presentation.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False))


def chart_analysis(presentation: Presentation) -> tuple[bool, dict]:
    charts = [shape.chart for slide in presentation.slides for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.CHART]
    matches = []
    for chart in charts:
        try:
            categories = [normalize(category.label) for category in chart.plots[0].categories]
            series = list(chart.series)
            values = [float(value) for value in series[0].values] if len(series) == 1 else []
            if categories == ["Complete", "Incomplete"] and values == [4.0, 3.0]:
                matches.append((categories, values, str(chart.chart_type)))
        except Exception:
            continue
    return len(matches) == 1, {"chart_count": len(charts), "matches": matches}


def table_analysis(presentation: Presentation, expected: dict) -> tuple[bool, dict]:
    candidates = []
    for slide_number, slide in enumerate(presentation.slides, 1):
        for shape in slide.shapes:
            if not shape.has_table:
                continue
            table = shape.table
            if len(table.rows) < 2:
                continue
            header = [normalize(cell.text).upper() for cell in table.rows[0].cells]
            required = {"ITEM", "OWNER", "STATUS", "DUE"}
            if not required.issubset(header):
                continue
            positions = {name: header.index(name) for name in required}
            observed = {}
            for row_index in range(1, len(table.rows)):
                row = table.rows[row_index]
                item = normalize(row.cells[positions["ITEM"]].text)
                if not item:
                    continue
                observed[item] = (
                    normalize(row.cells[positions["OWNER"]].text),
                    normalize(row.cells[positions["STATUS"]].text),
                    normalize(row.cells[positions["DUE"]].text),
                )
            candidates.append({"slide": slide_number, "header": header, "observed": observed})
    matches = [candidate for candidate in candidates if candidate["observed"] == expected]
    detail = [{"slide": item["slide"], "rows": item["observed"], "header": item["header"]} for item in candidates]
    return len(matches) == 1, {"candidate_count": len(candidates), "matches": len(matches), "detail": detail}


def non_table_slide_text(slide) -> str:
    return "\n".join(shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False) and not shape.has_table)


def incomplete_analysis(presentation: Presentation) -> tuple[bool, dict]:
    slides = [non_table_slide_text(slide) for slide in presentation.slides]
    detail = {}
    ok = True
    for item, (owner, status, due, _) in INCOMPLETE.items():
        matching_slides = [index + 1 for index, text in enumerate(slides) if all(token in text for token in (item, owner, status, due))]
        detail[item] = matching_slides
        ok = ok and bool(matching_slides)
    note_slides = [index + 1 for index, text in enumerate(slides) if NOTE in text]
    detail["note_slides"] = note_slides
    return ok and bool(note_slides), detail


def centers_overlap_x(first, second) -> bool:
    left = max(first.left, second.left)
    right = min(first.left + first.width, second.left + second.width)
    return right > left or abs((first.left + first.width / 2) - (second.left + second.width / 2)) <= 914400


def due_classification(presentation: Presentation) -> tuple[bool, dict]:
    detail = {}
    ok = True
    for item, (_, _, _, expected_label) in INCOMPLETE.items():
        associations = []
        for slide_number, slide in enumerate(presentation.slides, 1):
            item_shapes = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False) and item in shape.text]
            label_shapes = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False) and normalize(shape.text).upper() == expected_label]
            for item_shape in item_shapes:
                for label_shape in label_shapes:
                    if centers_overlap_x(item_shape, label_shape):
                        associations.append(slide_number)
        detail[item] = associations
        ok = ok and bool(associations)
    return ok, detail


def media_members(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as archive:
        return sorted(name for name in archive.namelist() if name.startswith("ppt/media/") and not name.endswith("/"))


def render(path: Path, outdir: Path) -> tuple[int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "profile"; profile.mkdir()
    completed = subprocess.run(["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(path)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300, check=False)
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return 0, completed.stdout[-500:]
    info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=30, check=False).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    return (int(match.group(1)) if match else 0), completed.stdout[-500:]


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {"criterion_id": identifier, "score": 1.0 if passed else 0.0, "evidence": evidence, "method": "deterministic", "raw": {"description": description, "weight": weight, "type": kind, "passed": bool(passed)}}


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [item["criterion_id"] for item in criteria if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0 and item["raw"]["passed"])
    penalty = sum(abs(item["raw"]["weight"]) for item in criteria if item["raw"]["weight"] < 0 and not item["raw"]["passed"])
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [sha256(path)] if path.is_file() else [], "criterion_results": criteria, "reward": reward, "passed": reward >= 0.8 and not hurdles, "hurdle_failures": hurdles}


def evaluate(path: Path, artifact_dir: Path) -> dict:
    try:
        expected, context = load_expected()
        candidate = Presentation(path)
    except Exception as exc:
        return make_result(path, [result_item(identifier, False, f"输入或 PPTX 打开失败：{exc}") for identifier in SPECS])

    criteria = []
    package_ok, package_evidence = package_checks(path)
    ratio = candidate.slide_width / candidate.slide_height if candidate.slide_height else 0
    file_ok = package_ok and len(candidate.slides) == 5 and abs(ratio - 16 / 9) < 0.02
    criteria.append(result_item("R001", file_ok, f"页数={len(candidate.slides)}，比例={ratio:.4f}，包={json.dumps(package_evidence, ensure_ascii=False)}"))

    text = all_text(candidate)
    upper = normalize(text).upper()
    percentages = re.findall(r"\b\d+(?:\.\d+)?%", text)
    summary_ok = "NOT READY" in upper and percentages and set(percentages) == {"57.1%"} and ("4 OF 7" in upper or "4 / 7" in upper) and ("3" in text and "INCOMPLETE" in upper or "OPEN / BLOCKED" in upper)
    criteria.append(result_item("R002", summary_ok, f"Not Ready={'NOT READY' in upper}，百分比={percentages}，4/7={('4 OF 7' in upper or '4 / 7' in upper)}，3未完成={('INCOMPLETE' in upper or 'OPEN / BLOCKED' in upper)}"))

    chart_ok, chart_detail = chart_analysis(candidate)
    criteria.append(result_item("R003", chart_ok, json.dumps(chart_detail, ensure_ascii=False)))

    table_ok, table_detail = table_analysis(candidate, expected)
    criteria.append(result_item("R004", table_ok, json.dumps(table_detail, ensure_ascii=False)))

    incomplete_ok, incomplete_detail = incomplete_analysis(candidate)
    criteria.append(result_item("R005", incomplete_ok, json.dumps(incomplete_detail, ensure_ascii=False)))

    due_ok, due_detail = due_classification(candidate)
    criteria.append(result_item("R006", due_ok, json.dumps(due_detail, ensure_ascii=False)))

    recipient_ok = all(recipient in text for recipient in RECIPIENTS) and CHANNEL in text
    invented_amounts = re.findall(r"(?:[$£€]\s?\d[\d,]*(?:\.\d+)?|\bUSD\s+\d[\d,]*(?:\.\d+)?)", text, flags=re.IGNORECASE)
    audience_ok = recipient_ok and not invented_amounts and set(context["report_audience"]) == RECIPIENTS and context["source_systems"]["slack"]["channel_name"] == CHANNEL
    criteria.append(result_item("R007", audience_ok, f"受众={recipient_ok}，编造金额={invented_amounts}"))

    tables = sum(shape.has_table for slide in candidate.slides for shape in slide.shapes)
    charts = sum(shape.shape_type == MSO_SHAPE_TYPE.CHART for slide in candidate.slides for shape in slide.shapes)
    full_slide_pictures = [shape.shape_id for slide in candidate.slides for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.width * shape.height >= candidate.slide_width * candidate.slide_height * 0.8]
    media = media_members(path)
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker in text.lower()]
    editable_ok = tables >= 1 and charts >= 1 and not full_slide_pictures and not media and not markers and not package_evidence.get("macros")
    criteria.append(result_item("R008", editable_ok, f"表格={tables}，图表={charts}，整页图={full_slide_pictures}，媒体={media}，标记={markers}"))

    pages, render_log = 0, ""
    try:
        with tempfile.TemporaryDirectory(prefix="automationbench4040-render-") as temporary:
            pages, render_log = render(path, Path(temporary))
    except Exception as exc:
        render_log = f"渲染异常：{exc}"
    delivered = sorted(str(item.relative_to(artifact_dir)) for item in artifact_dir.rglob("*") if item.is_file())
    hygiene_ok = pages == 5 and delivered == [OUTPUT_NAME] and package_ok
    criteria.append(result_item("R009", hygiene_ok, f"渲染页数={pages}，文件={delivered}，日志={render_log}"))

    severe = not (summary_ok and chart_ok and table_ok and incomplete_ok and due_ok and audience_ok and editable_ok)
    criteria.append(result_item("R010", not severe, f"严重错误={severe}"))
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    path = args.artifact_dir / OUTPUT_NAME
    result = evaluate(path, args.artifact_dir) if path.is_file() else make_result(path, [result_item(identifier, False, f"缺失指定交付物：{OUTPUT_NAME}") for identifier in SPECS])
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{result['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
