#!/usr/bin/env python3
from __future__ import annotations
import argparse,shutil,zipfile
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

OUTPUT_NAME="Three_Way_Match_Exception_Briefing.pptx"
CHART_NS="http://schemas.openxmlformats.org/drawingml/2006/chart"
def output(root,name):
    path=root/name/OUTPUT_NAME; path.parent.mkdir(parents=True,exist_ok=True); return path
def mutate(golden,target,operation):
    prs=Presentation(golden); operation(prs); prs.save(target)
def table(prs,header,columns):
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_table and len(shape.table.columns)==columns and shape.table.cell(0,0).text==header: return shape.table
    raise ValueError(header)
def replace(prs,old,new):
    count=0
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape,"has_text_frame",False) and old in shape.text:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if old in run.text: run.text=run.text.replace(old,new); count+=1
    if not count: raise ValueError(old)
def main():
    p=argparse.ArgumentParser(); p.add_argument("--golden",type=Path,required=True); p.add_argument("--outdir",type=Path,required=True); a=p.parse_args(); (a.outdir/"empty").mkdir(parents=True,exist_ok=True); output(a.outdir,"corrupt").write_bytes(b"bad\n")
    with zipfile.ZipFile(a.golden) as source,zipfile.ZipFile(output(a.outdir,"valid-repackage"),"w",compression=zipfile.ZIP_STORED) as target:
        for info in source.infolist(): target.writestr(info.filename,source.read(info.filename))
    def alt_cover(prs): prs.slides[0].background.fill.solid(); prs.slides[0].background.fill.fore_color.rgb=RGBColor(0x0D,0x2F,0x3A)
    mutate(a.golden,output(a.outdir,"valid-alt-cover"),alt_cover)
    mutate(a.golden,output(a.outdir,"wrong-source-row"),lambda prs:setattr(table(prs,"Invoice",9).cell(3,4),"text","30"))
    mutate(a.golden,output(a.outdir,"wrong-vi001-boundary"),lambda prs:setattr(table(prs,"Invoice",10).cell(1,8),"text","+2.10%"))
    mutate(a.golden,output(a.outdir,"wrong-vi002-status"),lambda prs:setattr(table(prs,"Invoice",10).cell(2,9),"text","Approved for Payment"))
    mutate(a.golden,output(a.outdir,"wrong-vi003-quantity"),lambda prs:setattr(table(prs,"Invoice",10).cell(3,2),"text","25"))
    mutate(a.golden,output(a.outdir,"wrong-vi003-status"),lambda prs:setattr(table(prs,"Invoice",10).cell(3,9),"text","Price Variance - Review"))
    def wrong_chart(prs):
        chart=next(shape.chart for slide in prs.slides for shape in slide.shapes if shape.shape_type==MSO_SHAPE_TYPE.CHART); caches=chart._chartSpace.findall(f".//{{{CHART_NS}}}numCache"); caches[1].findall(f"{{{CHART_NS}}}pt")[1].find(f"{{{CHART_NS}}}v").text="3500"
    mutate(a.golden,output(a.outdir,"wrong-chart"),wrong_chart)
    mutate(a.golden,output(a.outdir,"missing-policy-gap"),lambda prs:replace(prs,"policy only specifies “within 5%”","status accepted without caveat"))
    mutate(a.golden,output(a.outdir,"wrong-recipient"),lambda prs:replace(prs,"ap-lead@company.example.com","ap@company.example.com"))
    def contradiction(prs): box=prs.slides[-1].shapes.add_textbox(Inches(.7),Inches(6.55),Inches(3.4),Inches(.25)); box.text="VI-002 Approved for Payment"
    mutate(a.golden,output(a.outdir,"contradictory-approval"),contradiction)
    mutate(a.golden,output(a.outdir,"missing-citations"),lambda prs:replace(prs,"Source:","Sources omitted:"))
    image=a.outdir/"flat.png"; Image.new("RGB",(1600,900),"white").save(image)
    def flattened(prs):
        slide=prs.slides[2]
        for shape in list(slide.shapes): shape._element.getparent().remove(shape._element)
        slide.shapes.add_picture(str(image),0,0,width=prs.slide_width,height=prs.slide_height)
    mutate(a.golden,output(a.outdir,"flattened-decisions"),flattened)
    def marker(prs): box=prs.slides[-1].shapes.add_textbox(Inches(.1),Inches(.1),Inches(2),Inches(.2)); box.text="solution/Golden"
    mutate(a.golden,output(a.outdir,"hardcoded-marker"),marker)
    extra=a.outdir/"extra-output"; extra.mkdir(); shutil.copy2(a.golden,extra/OUTPUT_NAME); (extra/"debug.txt").write_text("debug\n")
if __name__=="__main__": main()
