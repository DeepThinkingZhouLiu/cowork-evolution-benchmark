#!/usr/bin/env python3
"""AutomationBench 4044 三单匹配简报任务专属 Judge。"""
from __future__ import annotations
import argparse,hashlib,json,posixpath,re,subprocess,tempfile,zipfile
from pathlib import Path
from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

OUTPUT_NAME="Three_Way_Match_Exception_Briefing.pptx"
TASK_ID="cowork-evo/automationbench-4044-three-way-match-exception-briefing"
RUBRIC_VERSION="automationbench-4044-ppt-adaptation-v0.1"
JUDGE_VERSION="automationbench-4044-deterministic-v0.1"
HASHES={"AutomationBench_4044_Context.json":"6698be721fdb27bbb35f2ab5d75960582bd6d825cff5453c1ebd7498f37e9a00","Three_Way_Match_Batch.xlsx":"03b91916d397ec4f11455650635543f2706bb044417b1ebdc7df3d2e5861bf98","Three_Way_Match_Tolerance_Policy.eml":"0413559546d2d14ba17e2fe4325bc98a312d2a736be5d84bff8d713df563adfc"}
SOURCE_HEADERS=["Invoice","Vendor","PO #","Inv qty","Received","PO qty","Invoice amount","PO amount","Receipt date"]
SOURCE_ROWS={
    "VI-001":("Acme Supplies","PO-5501","100","100","100","$5,100.00","$5,000.00","2026-02-05"),
    "VI-002":("Metro Supply","PO-5502","50","50","50","$3,750.00","$3,500.00","2026-02-06"),
    "VI-003":("TechWare Inc","PO-5503","25","25","30","$8,200.00","$9,600.00","2026-02-07"),
    "VI-004":("Global Logistics","PO-5504","200","200","200","$4,400.00","$4,400.00","2026-02-08"),
}
DECISION_HEADERS=["Invoice","Vendor","PO qty","Recv qty","Inv qty","PO amount","Inv amount","Δ amount","Δ %","Final status"]
DECISION_ROWS={
    "VI-001":("Acme Supplies","100","100","100","$5,000.00","$5,100.00","+$100.00","+2.00%","Approved for Payment"),
    "VI-002":("Metro Supply","50","50","50","$3,500.00","$3,750.00","+$250.00","+7.14%","Price Variance - Review"),
    "VI-003":("TechWare Inc","30","25","25","$9,600.00","$8,200.00","−$1,400.00","−14.58%","Quantity Discrepancy - Hold"),
    "VI-004":("Global Logistics","200","200","200","$4,400.00","$4,400.00","$0.00","0.00%","Approved for Payment"),
}
REQUIRED_SUMMARY=["$21,450","$22,500","$9,500","$11,950","50%","4 invoices reviewed"]
POLICY_GAP=["7.14%","policy does not define >5% price-only status","policy only specifies “within 5%”","Define treatment for price-only variance above 5%","upstream asserted Price Variance status"]
CONTROL=["quantity rule takes precedence","ap-lead@company.example.com","VI-002","VI-003","NON-APPROVED NOTICE"]
SUSPICIOUS=["solution/golden","judge-result","artifact_dir","automationbench-4044","/root/"]
SPECS={"R001":(6,"hurdle"),"R002":(16,"hurdle"),"R003":(24,"hurdle"),"R004":(16,"hurdle"),"R005":(12,"hurdle"),"R006":(12,"hurdle"),"R007":(4,"hurdle"),"R008":(8,"hurdle"),"R009":(2,"positive"),"R010":(-20,"penalty")}

def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def norm(value): return re.sub(r"\s+"," ",str(value).replace("\u00a0"," ").replace("\v"," ")).strip()
def locate(name):
    for p in (Path("/app/input_files")/name,Path(__file__).resolve().parents[1]/"data/input_files"/name):
        if p.is_file(): return p
    return None
def source_ok():
    missing=[]
    for name,digest in HASHES.items():
        path=locate(name)
        if path is None or sha(path)!=digest: missing.append(name)
    if not missing:
        wb=load_workbook(locate("Three_Way_Match_Batch.xlsx"),data_only=True)
        expected={"Invoice Batch":(5,7),"Purchase Orders":(5,5),"Receiving Reports":(5,4)}
        for sheet,shape in expected.items():
            if sheet not in wb.sheetnames or (wb[sheet].max_row,wb[sheet].max_column)!=shape: missing.append(f"{sheet}-shape")
        ctx=json.loads(locate("AutomationBench_4044_Context.json").read_text())
        if ctx.get("notify")!="ap-lead@company.example.com": missing.append("context-recipient")
    return not missing,missing
def package(path):
    evidence={"broken":None,"missing":[],"macros":[],"external":[]}
    try:
        with zipfile.ZipFile(path) as z:
            evidence["broken"]=z.testzip(); names=set(z.namelist()); evidence["macros"]=[n for n in names if n.lower().endswith("vbaproject.bin")]
            for relation in sorted(n for n in names if n.endswith(".rels")):
                root=etree.fromstring(z.read(relation)); owner="" if relation=="_rels/.rels" else posixpath.join(relation.split("/_rels/",1)[0],relation.split("/_rels/",1)[1][:-5]); base=posixpath.dirname(owner)
                for rel in root:
                    target=rel.attrib.get("Target","")
                    if rel.attrib.get("TargetMode")=="External": evidence["external"].append(target); continue
                    resolved=target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base,target))
                    if resolved not in names: evidence["missing"].append(f"{relation}->{resolved}")
        return evidence["broken"] is None and not evidence["missing"] and not evidence["macros"],evidence
    except Exception as exc: evidence["exception"]=str(exc); return False,evidence
def deck_text(prs): return "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if getattr(shape,"has_text_frame",False))
def tables(prs):
    out=[]
    for slide_no,slide in enumerate(prs.slides,1):
        for shape in slide.shapes:
            if shape.has_table: out.append((slide_no,[[norm(c.text) for c in row.cells] for row in shape.table.rows]))
    return out
def table_map(all_tables,headers): return [(n,{row[0]:tuple(row[1:]) for row in rows[1:] if row and row[0]}) for n,rows in all_tables if rows and rows[0]==headers]
def charts(prs):
    out=[]
    for slide_no,slide in enumerate(prs.slides,1):
        for shape in slide.shapes:
            if shape.shape_type==MSO_SHAPE_TYPE.CHART:
                chart=shape.chart; cats=[norm(c.label) for c in chart.plots[0].categories]; series={norm(s.name):[float(v) for v in s.values] for s in chart.series}; out.append((slide_no,cats,series))
    return out
def render(path,out):
    out.mkdir(parents=True,exist_ok=True); profile=out/"profile"; profile.mkdir(exist_ok=True)
    proc=subprocess.run(["soffice",f"-env:UserInstallation=file://{profile}","--headless","--convert-to","pdf","--outdir",str(out),str(path)],stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,timeout=300); pdf=out/f"{path.stem}.pdf"
    if proc.returncode or not pdf.is_file(): return 0,proc.stdout[-500:]
    info=subprocess.run(["pdfinfo",str(pdf)],stdout=subprocess.PIPE,text=True).stdout; match=re.search(r"^Pages:\s+(\d+)",info,re.M); return int(match.group(1)) if match else 0,proc.stdout[-500:]
def item(cid,passed,evidence):
    weight,kind=SPECS[cid]; return {"criterion_id":cid,"score":1.0 if passed else 0.0,"evidence":evidence,"method":"deterministic","raw":{"weight":weight,"type":kind,"passed":bool(passed)}}
def assemble(path,criteria):
    hurdles=[x["criterion_id"] for x in criteria if x["raw"]["type"]=="hurdle" and not x["raw"]["passed"]]; total=sum(x["raw"]["weight"] for x in criteria if x["raw"]["weight"]>0); earned=sum(x["raw"]["weight"] for x in criteria if x["raw"]["weight"]>0 and x["raw"]["passed"]); penalty=sum(abs(x["raw"]["weight"]) for x in criteria if x["raw"]["weight"]<0 and not x["raw"]["passed"]); reward=0 if hurdles else max(0,min(1,(earned-penalty)/total)); return {"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":[sha(path)] if path.is_file() else [],"criterion_results":criteria,"reward":reward,"passed":reward>=.8 and not hurdles,"hurdle_failures":hurdles}
def evaluate(path,artifact):
    src,src_evidence=source_ok()
    try: prs=Presentation(path)
    except Exception as exc: return assemble(path,[item(cid,False,f"打开失败:{exc}") for cid in SPECS])
    pkg,pkg_evidence=package(path); ratio=prs.slide_width/prs.slide_height if prs.slide_height else 0; r1=src and pkg and len(prs.slides)==6 and abs(ratio-16/9)<.02; criteria=[item("R001",r1,f"pages={len(prs.slides)} ratio={ratio} source={src_evidence} package={json.dumps(pkg_evidence,ensure_ascii=False)}")]
    all_tables=tables(prs); source=table_map(all_tables,SOURCE_HEADERS); r2=len(source)==1 and source[0][1]==SOURCE_ROWS; criteria.append(item("R002",r2,f"source={source}"))
    decisions=table_map(all_tables,DECISION_HEADERS); r3=len(decisions)==1 and decisions[0][1]==DECISION_ROWS; criteria.append(item("R003",r3,f"decisions={decisions}"))
    text=deck_text(prs); chart_data=charts(prs); expected_chart=(['VI-001','VI-002','VI-003','VI-004'],{'PO amount':[5000.0,3500.0,9600.0,4400.0],'Invoice amount':[5100.0,3750.0,8200.0,4400.0]}); chart_ok=sum(1 for _,cats,series in chart_data if (cats,series)==expected_chart)==1; missing_summary=[x for x in REQUIRED_SUMMARY if x not in text]; r4=chart_ok and not missing_summary; criteria.append(item("R004",r4,f"missing_summary={missing_summary} charts={chart_data}"))
    missing_gap=[x for x in POLICY_GAP if x.lower() not in text.lower()]; false_claims=[x for x in ["policy authorizes >5%","policy covers 7.14%","within policy at 7.14%"] if x in text.lower()]; r5=not missing_gap and not false_claims; criteria.append(item("R005",r5,f"missing={missing_gap} false_claims={false_claims}"))
    missing_control=[x for x in CONTROL if x.lower() not in text.lower()]; contradictions=[x for x in ["VI-002 Approved for Payment","VI-003 Approved for Payment"] if x.lower() in text.lower()]; r6=not missing_control and not contradictions and text.lower().count("ap-lead@company.example.com")>=2; criteria.append(item("R006",r6,f"missing={missing_control} contradictions={contradictions} recipient_count={text.lower().count('ap-lead@company.example.com')}") )
    citations=text.count("Source:"); tags=text.count("AutomationBench 4044"); markers=[x for x in SUSPICIOUS if x in text.lower()]; r7=citations==6 and tags==6 and not markers; criteria.append(item("R007",r7,f"citations={citations} source_tags={tags} markers={markers}"))
    native_tables=sum(shape.has_table for slide in prs.slides for shape in slide.shapes); native_charts=sum(shape.shape_type==MSO_SHAPE_TYPE.CHART for slide in prs.slides for shape in slide.shapes); full_page=[shape.shape_id for slide in prs.slides for shape in slide.shapes if shape.shape_type==MSO_SHAPE_TYPE.PICTURE and shape.width*shape.height>=prs.slide_width*prs.slide_height*.8]
    with zipfile.ZipFile(path) as z: media=[n for n in z.namelist() if n.startswith("ppt/media/") and not n.endswith("/")]
    r8=native_tables>=2 and native_charts>=1 and not full_page and not media and not pkg_evidence.get("macros") and not pkg_evidence.get("external"); criteria.append(item("R008",r8,f"tables={native_tables} charts={native_charts} full_page={full_page} media={media} external={pkg_evidence.get('external')}"))
    pages=0; log=""
    try:
        with tempfile.TemporaryDirectory(prefix="ab4044-render-") as directory: pages,log=render(path,Path(directory))
    except Exception as exc: log=str(exc)
    delivered=sorted(str(p.relative_to(artifact)) for p in artifact.rglob("*") if p.is_file()); r9=pages==6 and delivered==[OUTPUT_NAME] and pkg; criteria.append(item("R009",r9,f"pages={pages} files={delivered} log={log}"))
    severe=not all([r2,r3,r4,r5,r6,r7,r8]); criteria.append(item("R010",not severe,f"severe={severe}")); return assemble(path,criteria)
def main():
    p=argparse.ArgumentParser(); p.add_argument("--artifact-dir",type=Path,required=True); p.add_argument("--result",type=Path,required=True); p.add_argument("--reward-file",type=Path,required=True); a=p.parse_args(); path=a.artifact_dir/OUTPUT_NAME; result=evaluate(path,a.artifact_dir) if path.is_file() else assemble(path,[item(cid,False,f"缺失 {OUTPUT_NAME}") for cid in SPECS]); a.result.parent.mkdir(parents=True,exist_ok=True); a.reward_file.parent.mkdir(parents=True,exist_ok=True); a.result.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n"); a.reward_file.write_text(f"{result['reward']:.6f}\n"); print(json.dumps(result,ensure_ascii=False,indent=2))
if __name__=="__main__": main()
