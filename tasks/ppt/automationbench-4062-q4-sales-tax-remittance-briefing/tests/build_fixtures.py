#!/usr/bin/env python3
from __future__ import annotations
import argparse,shutil,zipfile
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches
OUTPUT_NAME="Q4_2025_Sales_Tax_Remittance_Briefing.pptx"; C="http://schemas.openxmlformats.org/drawingml/2006/chart"
def out(root,name): p=root/name/OUTPUT_NAME; p.parent.mkdir(parents=True,exist_ok=True); return p
def mutate(g,o,fn): p=Presentation(g); fn(p); p.save(o)
def table(prs,header):
    for sl in prs.slides:
        for s in sl.shapes:
            if s.has_table and s.table.cell(0,0).text==header: return s.table
    raise ValueError(header)
def replace(prs,old,new):
    n=0
    for sl in prs.slides:
        for s in sl.shapes:
            if getattr(s,"has_text_frame",False) and old in s.text: s.text=s.text.replace(old,new); n+=1
    if not n: raise ValueError(old)
def main():
    p=argparse.ArgumentParser(); p.add_argument("--golden",type=Path,required=True); p.add_argument("--outdir",type=Path,required=True); a=p.parse_args(); (a.outdir/"empty").mkdir(parents=True); out(a.outdir,"corrupt").write_bytes(b"bad\n")
    with zipfile.ZipFile(a.golden) as r,zipfile.ZipFile(out(a.outdir,"valid-repackage"),"w",compression=zipfile.ZIP_STORED) as w:
        for i in r.infolist(): w.writestr(i.filename,r.read(i.filename))
    def alt(prs): prs.slides[0].background.fill.solid(); prs.slides[0].background.fill.fore_color.rgb=RGBColor(0x10,0x2B,0x35)
    mutate(a.golden,out(a.outdir,"valid-alt-cover"),alt)
    def missing(prs):
        t=table(prs,"row_id")
        for c in t.rows[4].cells: c.text=""
    mutate(a.golden,out(a.outdir,"missing-source-row"),missing)
    mutate(a.golden,out(a.outdir,"wrong-source-sale"),lambda prs:setattr(table(prs,"row_id").cell(6,3),"text","$175,000"))
    mutate(a.golden,out(a.outdir,"wrong-ca-dec-rate"),lambda prs:setattr(table(prs,"State").cell(3,5),"text","7.25%"))
    mutate(a.golden,out(a.outdir,"wrong-ny-exemption"),lambda prs:setattr(table(prs,"State").cell(5,3),"text","≈$20,000"))
    mutate(a.golden,out(a.outdir,"wrong-month-tax"),lambda prs:setattr(table(prs,"State").cell(3,6),"text","$11,962.50"))
    def chart(prs):
        c=next(s.chart for sl in prs.slides for s in sl.shapes if s.shape_type==MSO_SHAPE_TYPE.CHART); cache=c._chartSpace.find(f".//{{{C}}}numCache"); cache.findall(f"{{{C}}}pt")[0].find(f"{{{C}}}v").text="31175"
    mutate(a.golden,out(a.outdir,"wrong-liability-chart"),chart)
    mutate(a.golden,out(a.outdir,"wrong-summary-total"),lambda prs:setattr(table(prs,"Jurisdiction").cell(4,4),"text","$47,975.00"))
    mutate(a.golden,out(a.outdir,"missing-approximation-caveat"),lambda prs:replace(prs,"asserted benchmark arithmetic","ordinary accounting estimate"))
    def filing(prs): b=prs.slides[-1].shapes.add_textbox(Inches(.8),Inches(6.2),Inches(2.5),Inches(.25)); b.text="FILING COMPLETE"
    mutate(a.golden,out(a.outdir,"filing-complete"),filing)
    def payment(prs): b=prs.slides[-1].shapes.add_textbox(Inches(3.6),Inches(6.2),Inches(2.5),Inches(.25)); b.text="PAYMENT INITIATED"
    mutate(a.golden,out(a.outdir,"payment-initiated"),payment)
    mutate(a.golden,out(a.outdir,"wrong-recipient"),lambda prs:replace(prs,"tax@company.example.com","accounting@company.example.com"))
    mutate(a.golden,out(a.outdir,"missing-citations"),lambda prs:replace(prs,"Source:","Sources omitted:"))
    image=a.outdir/"flat.png"; Image.new("RGB",(1600,900),"white").save(image)
    def flat(prs):
        sl=prs.slides[2]
        for s in list(sl.shapes): s._element.getparent().remove(s._element)
        sl.shapes.add_picture(str(image),0,0,width=prs.slide_width,height=prs.slide_height)
    mutate(a.golden,out(a.outdir,"flattened-schedule"),flat)
    def marker(prs): b=prs.slides[-1].shapes.add_textbox(Inches(.1),Inches(.1),Inches(2),Inches(.2)); b.text="solution/Golden"
    mutate(a.golden,out(a.outdir,"hardcoded-marker"),marker)
    extra=a.outdir/"extra-output"; extra.mkdir(); shutil.copy2(a.golden,extra/OUTPUT_NAME); (extra/"debug.txt").write_text("debug\n")
if __name__=="__main__": main()
