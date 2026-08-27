#!/usr/bin/env python3
"""Cowork-Bench ``sf-department-budget-ppt`` 的单题可回放参考解。

脚本只处理原题两个上游输入：部门预算 PDF 与原始 PostgreSQL COPY
快照。所有数字都在运行时从输入重算，不包含题目真值常量。PPTX 由
任务内随包的 MIT 许可 presentation-skill 渲染器根据 outline.json 构建。
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


PDF_NAME = "department_budgets_fy2026.pdf"
EMPLOYEES_NAME = "HR_ANALYTICS__PUBLIC__EMPLOYEES.copy.sql"
XLSX_NAME = "budget_vs_actuals.xlsx"
PPTX_NAME = "budget_vs_actuals.pptx"

DETAIL_HEADERS = [
    "Department",
    "Approved_Budget",
    "Actual_Expenditure",
    "Variance",
    "Variance_Pct",
    "Headcount",
    "Avg_Salary",
]
TOTAL_LABELS = [
    "Total_Departments",
    "Total_Budget",
    "Total_Actuals",
    "Total_Variance",
    "Most_Under_Budget",
    "Most_Over_Budget",
]

NAVY = "0F172A"
BLUE = "1D4ED8"
TEAL = "0891B2"
PALE = "F8FAFC"
WHITE = "FFFFFF"
INK = "111827"
MUTED = "4B5563"
LINE = "D1D5DB"


def money(value: Decimal) -> str:
    return f"${value:,.2f}"


def pct(value: Decimal) -> str:
    return f"{value:,.2f}%"


def q2(value: Decimal) -> Decimal:
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def run(command: list[str], *, timeout: int = 240) -> subprocess.CompletedProcess[str]:
    completed = subprocess.run(
        command,
        capture_output=True,
        text=True,
        timeout=timeout,
        check=False,
    )
    if completed.returncode != 0:
        detail = (completed.stderr or completed.stdout or "命令执行失败")[-2000:]
        raise RuntimeError(f"{' '.join(command[:3])} 失败：{detail}")
    return completed


def read_budgets(pdf_path: Path) -> dict[str, Decimal]:
    executable = shutil.which("pdftotext")
    if executable is None:
        raise RuntimeError("缺少 pdftotext，无法从上游 PDF 读取预算")
    text = run([executable, "-layout", str(pdf_path), "-"]).stdout
    budgets: dict[str, Decimal] = {}
    row_pattern = re.compile(
        r"^(?P<department>\S(?:.*?\S)?)\s+\$(?P<amount>[0-9][0-9,]*\.\d{2})$"
    )
    for raw_line in text.splitlines():
        match = row_pattern.match(raw_line.strip())
        if not match:
            continue
        department = match.group("department")
        if department == "Итого":
            continue
        budgets[department] = Decimal(match.group("amount").replace(",", ""))
    if len(budgets) != 7:
        raise ValueError(f"预算 PDF 应包含 7 个部门，实际为 {sorted(budgets)}")
    return budgets


def parse_copy_header(line: str) -> list[str]:
    if not line.startswith('COPY sf_data."HR_ANALYTICS__PUBLIC__EMPLOYEES" '):
        raise ValueError("输入不是原题指定的 EMPLOYEES COPY 块")
    columns = re.findall(r'"([^"]+)"', line)
    if columns and columns[0] == "HR_ANALYTICS__PUBLIC__EMPLOYEES":
        columns = columns[1:]
    if len(columns) != 18:
        raise ValueError(f"COPY 表头字段数异常：{columns}")
    return columns


def read_actuals(copy_path: Path) -> dict[str, dict[str, Decimal | int]]:
    aggregates: dict[str, dict[str, Decimal | int]] = defaultdict(
        lambda: {"headcount": 0, "actual": Decimal("0")}
    )
    with copy_path.open(encoding="utf-8", newline="") as stream:
        header = stream.readline().rstrip("\r\n")
        columns = parse_copy_header(header)
        department_index = columns.index("DEPARTMENT")
        salary_index = columns.index("SALARY")
        row_count = 0
        for raw_line in stream:
            line = raw_line.rstrip("\r\n")
            if line == "\\.":
                break
            fields = line.split("\t")
            if len(fields) != len(columns):
                raise ValueError(f"第 {row_count + 1} 条员工记录字段数不完整")
            department = fields[department_index]
            salary = Decimal(fields[salary_index])
            aggregates[department]["headcount"] = int(aggregates[department]["headcount"]) + 1
            aggregates[department]["actual"] = Decimal(aggregates[department]["actual"]) + salary
            row_count += 1
        else:
            raise ValueError("COPY 快照缺少 \\. 终止行")
    if row_count == 0:
        raise ValueError("COPY 快照中没有员工记录")
    return aggregates


def calculate(input_dir: Path) -> tuple[list[dict[str, Decimal | int | str]], dict[str, Decimal | int | str]]:
    budgets = read_budgets(input_dir / PDF_NAME)
    actuals = read_actuals(input_dir / EMPLOYEES_NAME)
    if set(budgets) != set(actuals):
        raise ValueError(
            f"预算与员工表的部门集不一致："
            f"PDF={sorted(budgets)} COPY={sorted(actuals)}"
        )

    rows: list[dict[str, Decimal | int | str]] = []
    for department in sorted(budgets):
        budget = q2(budgets[department])
        actual = q2(Decimal(actuals[department]["actual"]))
        headcount = int(actuals[department]["headcount"])
        variance = q2(budget - actual)
        variance_pct = q2(variance / budget * 100)
        avg_salary = q2(actual / headcount)
        rows.append(
            {
                "department": department,
                "budget": budget,
                "actual": actual,
                "variance": variance,
                "variance_pct": variance_pct,
                "headcount": headcount,
                "avg_salary": avg_salary,
                "status": "Under Budget" if variance >= 0 else "Over Budget",
            }
        )

    total_budget = q2(sum((Decimal(row["budget"]) for row in rows), Decimal("0")))
    total_actual = q2(sum((Decimal(row["actual"]) for row in rows), Decimal("0")))
    total_variance = q2(total_budget - total_actual)
    most_under = max(rows, key=lambda row: Decimal(row["variance"]))
    most_over = min(rows, key=lambda row: Decimal(row["variance"]))
    over_count = sum(1 for row in rows if Decimal(row["variance"]) < 0)
    totals: dict[str, Decimal | int | str] = {
        "total_departments": len(rows),
        "total_budget": total_budget,
        "total_actual": total_actual,
        "total_variance": total_variance,
        "most_under": str(most_under["department"]),
        "most_under_variance": Decimal(most_under["variance"]),
        "most_over": str(most_over["department"]),
        "most_over_variance": Decimal(most_over["variance"]),
        "over_count": over_count,
    }
    return rows, totals


def style_table(ws, widths: list[float]) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 28
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", size=11, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color=LINE)
    for row_index in range(2, ws.max_row + 1):
        ws.row_dimensions[row_index].height = 24
        for cell in ws[row_index]:
            cell.fill = PatternFill("solid", fgColor=WHITE if row_index % 2 == 0 else PALE)
            cell.font = Font(name="Aptos", size=10, color=INK)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(
                horizontal="left" if cell.column == 1 else "right",
                vertical="center",
            )
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + index)].width = width


def configure_print_layout(ws, *, landscape: bool, print_area: str) -> None:
    """让 LibreOffice/PDF 打印版保持一张表在单页宽度内。"""

    ws.print_area = print_area
    ws.print_title_rows = "1:1"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4
    ws.page_margins.header = 0.15
    ws.page_margins.footer = 0.15


def build_workbook(rows: list[dict[str, Decimal | int | str]], totals: dict[str, Decimal | int | str], path: Path) -> None:
    workbook = Workbook()
    details = workbook.active
    details.title = "Department Details"
    details.append(DETAIL_HEADERS)
    for index, row in enumerate(rows, start=2):
        details.append(
            [
                row["department"],
                float(Decimal(row["budget"])),
                float(Decimal(row["actual"])),
                float(Decimal(row["variance"])),
                float(Decimal(row["variance_pct"])),
                int(row["headcount"]),
                float(Decimal(row["avg_salary"])),
            ]
        )
    style_table(details, [20, 21, 22, 20, 18, 14, 16])
    table = Table(displayName="DepartmentDetailsTable", ref=f"A1:G{details.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    details.add_table(table)
    for row_index in range(2, details.max_row + 1):
        for column in ("B", "C", "D", "G"):
            details[f"{column}{row_index}"].number_format = '$#,##0.00;[Red]-$#,##0.00'
        details[f"E{row_index}"].number_format = '0.00"%";[Red]-0.00"%"'
        details[f"F{row_index}"].number_format = "#,##0"
    configure_print_layout(details, landscape=True, print_area=f"A1:G{details.max_row}")

    total_sheet = workbook.create_sheet("Totals")
    total_sheet.append(["Metric", "Value"])
    total_values = [
        totals["total_departments"],
        float(Decimal(totals["total_budget"])),
        float(Decimal(totals["total_actual"])),
        float(Decimal(totals["total_variance"])),
        totals["most_under"],
        totals["most_over"],
    ]
    for label, value in zip(TOTAL_LABELS, total_values):
        total_sheet.append([label, value])
    style_table(total_sheet, [25, 24])
    totals_table = Table(displayName="TotalsTable", ref=f"A1:B{total_sheet.max_row}")
    totals_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    total_sheet.add_table(totals_table)
    for row_index in (3, 4, 5):
        total_sheet[f"B{row_index}"].number_format = '$#,##0.00;[Red]-$#,##0.00'
    configure_print_layout(total_sheet, landscape=False, print_area=f"A1:B{total_sheet.max_row}")

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


def recalculate_workbook(path: Path) -> None:
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if executable is None:
        raise RuntimeError("缺少 LibreOffice，无法重算 Excel 公式")
    with tempfile.TemporaryDirectory(prefix="cowork-evo-budget-xlsx-") as temp_dir:
        temp_root = Path(temp_dir)
        output_dir = temp_root / "converted"
        output_dir.mkdir()
        profile = temp_root / "profile"
        run(
            [
                executable,
                "--headless",
                f"-env:UserInstallation={profile.resolve().as_uri()}",
                "--convert-to",
                "xlsx",
                "--outdir",
                str(output_dir),
                str(path),
            ]
        )
        recalculated = output_dir / path.name
        if not recalculated.is_file():
            raise RuntimeError("LibreOffice 未生成重算后的 XLSX")
        shutil.copy2(recalculated, path)


def department_slide(row: dict[str, Decimal | int | str], index: int) -> dict:
    department = str(row["department"])
    common = {
        "type": "content",
        "title": department,
        "subtitle": f"FY2026 budget vs salary actuals | {row['status']}",
        "slide_intent": "evidence",
        "visual_intent": "comparison",
        "sources": [PDF_NAME, EMPLOYEES_NAME],
    }
    # 上游用户明确要求“每个部门一页”。同一种 Plan vs Actual
    # 两栏结构使管理层能够逐页横向对比，同时避免长金额在窄卡片中断行。
    return {
        **common,
        "variant": "comparison-2col",
        "left": {
            "title": "Approved Plan",
            "body": [
                f"Approved Budget: {money(Decimal(row['budget']))}",
                f"Headcount: {int(row['headcount']):,}",
                f"Avg Salary: {money(Decimal(row['avg_salary']))}",
            ],
        },
        "right": {
            "title": "Salary Actuals",
            "body": [
                f"Actual Expenditure: {money(Decimal(row['actual']))}",
                f"Variance: {money(Decimal(row['variance']))}",
                f"Variance Pct: {pct(Decimal(row['variance_pct']))}",
            ],
        },
        "verdict": f"Status: {row['status']}",
    }


def make_outline(rows: list[dict[str, Decimal | int | str]], totals: dict[str, Decimal | int | str]) -> dict:
    over_count = int(totals["over_count"])
    if over_count == 0:
        verdict = f"All {int(totals['total_departments'])} departments are under budget."
    elif over_count == 1:
        verdict = f"1 of {int(totals['total_departments'])} departments is over budget."
    else:
        verdict = f"{over_count} of {int(totals['total_departments'])} departments are over budget."

    slides: list[dict] = [
        {
            "type": "title",
            "title": "FY2026 Budget vs Actuals",
            "subtitle": "Quarterly Leadership Review - March 2026",
            "footer": "CFO leadership briefing | Source-backed salary actuals",
        },
        {
            "type": "content",
            "variant": "table",
            "title": "Department Overview",
            "subtitle": "Approved FY2026 budgets compared with salary actuals",
            "slide_intent": "evidence",
            "visual_intent": "data",
            "headers": ["Metric", "Value"],
            "rows": [
                ["Departments", f"{int(totals['total_departments'])}"],
                ["Total Budget", money(Decimal(totals["total_budget"]))],
                ["Total Actuals", money(Decimal(totals["total_actual"]))],
                ["Total Variance", money(Decimal(totals["total_variance"]))],
            ],
            "column_weights": [1.0, 1.5],
            "caption": "Source-backed roll-up across all seven departments",
            "sources": [PDF_NAME, EMPLOYEES_NAME],
        },
    ]
    slides.extend(department_slide(row, index) for index, row in enumerate(rows))
    row_by_department = {str(row["department"]): row for row in rows}
    most_under_row = row_by_department[str(totals["most_under"])]
    most_over_row = row_by_department[str(totals["most_over"])]
    slides.append(
        {
            "type": "content",
            "variant": "comparison-2col",
            "title": "Summary",
            "subtitle": "Variance extremes for leadership review",
            "slide_intent": "decision",
            "visual_intent": "comparison",
            "left": {
                "title": "Most Under Budget",
                "body": [
                    str(totals["most_under"]),
                    f"Approved Budget: {money(Decimal(most_under_row['budget']))}",
                    f"Actual Expenditure: {money(Decimal(most_under_row['actual']))}",
                    f"Variance: {money(Decimal(totals['most_under_variance']))}",
                ],
            },
            "right": {
                "title": "Most Over Budget",
                "body": [
                    str(totals["most_over"]),
                    f"Approved Budget: {money(Decimal(most_over_row['budget']))}",
                    f"Actual Expenditure: {money(Decimal(most_over_row['actual']))}",
                    f"Variance: {money(Decimal(totals['most_over_variance']))}",
                ],
            },
            "verdict": verdict,
            "sources": [PDF_NAME, EMPLOYEES_NAME],
        }
    )
    return {
        "title": "FY2026 Budget vs Actuals",
        "subtitle": "Quarterly Leadership Review - March 2026",
        "deck_style": {
            "font_pair": "clean_modern_v1",
            "palette_key": "enterprise_graphite_v1",
            "visual_density": "medium",
            "emoji_mode": "none",
            "title_motif": "geometric",
            "show_page_numbers": False,
        },
        "compliance": {"require_attribution": False},
        "slides": slides,
    }


def write_workspace(
    workspace_dir: Path,
    outline: dict,
    rows: list[dict[str, Decimal | int | str]],
    totals: dict[str, Decimal | int | str],
) -> None:
    workspace_dir.mkdir(parents=True, exist_ok=True)
    (workspace_dir / "outline.json").write_text(
        json.dumps(outline, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    truth = {"departments": rows, "totals": totals}
    (workspace_dir / "truth.json").write_text(
        json.dumps(truth, ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )


def build_deck(workspace_dir: Path, output_path: Path, builder_path: Path) -> None:
    run(
        [
            sys.executable,
            str(builder_path),
            "--outline",
            str(workspace_dir / "outline.json"),
            "--output",
            str(output_path),
            "--style-preset",
            "data-heavy-boardroom",
            "--layout-16x9",
            "--overwrite",
        ],
        timeout=300,
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=Path("/app/data/input_files"))
    parser.add_argument("--output-dir", type=Path, default=Path("/app/output"))
    parser.add_argument("--workspace-dir", type=Path)
    parser.add_argument("--builder", type=Path)
    args = parser.parse_args()

    script_dir = Path(__file__).resolve().parent
    builder = args.builder or script_dir / "vendor" / "presentation_skill" / "build_deck.py"
    workspace_dir = args.workspace_dir or args.output_dir / ".deck_workspace"
    for required in (args.input_dir / PDF_NAME, args.input_dir / EMPLOYEES_NAME, builder):
        if not required.is_file():
            raise FileNotFoundError(f"缺少必需文件：{required}")

    rows, totals = calculate(args.input_dir)
    outline = make_outline(rows, totals)
    write_workspace(workspace_dir, outline, rows, totals)
    args.output_dir.mkdir(parents=True, exist_ok=True)

    workbook_path = args.output_dir / XLSX_NAME
    build_workbook(rows, totals, workbook_path)
    build_deck(workspace_dir, args.output_dir / PPTX_NAME, builder)

    print(json.dumps({"departments": rows, "totals": totals}, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
