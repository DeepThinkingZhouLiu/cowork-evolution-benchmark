#!/usr/bin/env python3
"""Cowork-Bench ppt-clickhouse-executive 离线改造题 Judge。

评分时不读取 Golden，而是从两个上游静态快照重新连接、聚合并计算
Q4 2024 真值，然后分别检查 Excel、PPTX、PDF 以及三者之间的一致性。
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import subprocess
import tempfile
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


TASK_ID = "cowork-evo/cowork-bench-q4-executive-briefing-package"
RUBRIC_VERSION = "cowork-bench-ppt-clickhouse-executive-adaptation-v0.1"
JUDGE_VERSION = "cowork-bench-ppt-clickhouse-executive-offline-v0.1"

ORDERS_NAME = "q4_orders_2024.csv"
CUSTOMERS_NAME = "customer_dimension.csv"
XLSX_NAME = "Q4_Executive_Analysis.xlsx"
PPTX_NAME = "Q4_Executive_Briefing.pptx"
PDF_NAME = "Q4_Executive_Briefing.pdf"

ORDERS_SHA256 = "432161d192c7239e634350125d9f5347cdc4b0d70a7f264a463cbfb351fb4ce3"
CUSTOMERS_SHA256 = "0f0e68e115c96312a1fa0188912689271d6a15e0765f915fa15e078d7a34a575"

REGIONAL_HEADERS = ["Region", "Orders", "Revenue", "Avg_Order_Value"]
SEGMENT_HEADERS = ["Segment", "Orders", "Revenue", "Revenue_Share_Pct"]

CRITERION_SPECS = {
    "R001": {
        "description": "三个指定交付文件齐全且可打开；Excel 为 2 张表、PPTX 和 PDF 均为 7 页。",
        "weight": 4,
        "type": "hurdle",
        "cap": 0.0,
    },
    "R002": {
        "description": "Regional Summary 的表头、5 个地区、排序、订单、收入和客单价与输入重算真值一致。",
        "weight": 8,
        "type": "hurdle",
        "cap": 0.0,
    },
    "R003": {
        "description": "Segment Analysis 的表头、4 个客户分群、排序、订单、收入和收入占比与输入重算真值一致。",
        "weight": 7,
        "type": "hurdle",
        "cap": 0.0,
    },
    "R004": {
        "description": "Excel 具备专业可用的表格样式、冻结表头、筛选、列宽和数字格式。",
        "weight": 3,
        "type": "positive",
        "cap": 1.0,
    },
    "R005": {
        "description": "PPTX 严格按标题页 -> 5 个字母序地区页 -> 汇总页组织，标题准确。",
        "weight": 5,
        "type": "hurdle",
        "cap": 0.4,
    },
    "R006": {
        "description": "5 个地区页均显示对应地区的订单数和收入，与输入真值一致。",
        "weight": 6,
        "type": "hurdle",
        "cap": 0.5,
    },
    "R007": {
        "description": "最后一页准确显示季度总订单数和总收入。",
        "weight": 5,
        "type": "hurdle",
        "cap": 0.5,
    },
    "R008": {
        "description": "PPTX 为 16:9，主要文字均在画布内，页面由可编辑文字和形状组成，不是整页压平图。",
        "weight": 4,
        "type": "positive",
        "cap": 1.0,
    },
    "R009": {
        "description": "PDF 归档为 7 页 16:9，按页保留标题、地区、地区指标和最终汇总。",
        "weight": 5,
        "type": "hurdle",
        "cap": 0.4,
    },
    "R010": {
        "description": "Excel、PPTX 和 PDF 中的地区指标、总订单和总收入交叉一致。",
        "weight": 5,
        "type": "positive",
        "cap": 1.0,
    },
    "R011": {
        "description": "PPTX 能由 LibreOffice 实际渲染为 7 页 PDF，关键按页内容仍可读。",
        "weight": 3,
        "type": "positive",
        "cap": 1.0,
    },
    "R012": {
        "description": "Excel 能由 LibreOffice 实际打开、重算并另存，两张分析表结构保持。",
        "weight": 2,
        "type": "positive",
        "cap": 1.0,
    },
    "R013": {
        "description": "若任一指定交付物缺失、损坏或 PDF 为空壳，扣 12 分。",
        "weight": -12,
        "type": "penalty",
        "cap": 1.0,
    },
    "R014": {
        "description": "若 PPTX 存在覆盖全页的大图或 PDF 与 PPTX 的核心按页内容不一致，扣 10 分。",
        "weight": -10,
        "type": "penalty",
        "cap": 1.0,
    },
}


def digest(path: Path) -> str:
    value = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            value.update(chunk)
    return value.hexdigest()


def quantize(value: Decimal, places: str) -> Decimal:
    return value.quantize(Decimal(places), rounding=ROUND_HALF_UP)


def normalize_text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold()


def numeric_values(text: str) -> list[Decimal]:
    """抽取英文或俄文展示中的数字。"""
    normalized = str(text).replace("\xa0", " ").replace("\u202f", " ").replace("\u2009", " ")
    tokens = re.findall(r"(?<![\w])[-+]?\d[\d ,.]*\d|(?<![\w])[-+]?\d", normalized)
    values: list[Decimal] = []
    for token in tokens:
        compact = token.strip().replace(" ", "")
        if not compact:
            continue
        if "," in compact and "." in compact:
            if compact.rfind(".") > compact.rfind(","):
                compact = compact.replace(",", "")
            else:
                compact = compact.replace(".", "").replace(",", ".")
        elif compact.count(",") == 1:
            tail = compact.split(",", 1)[1]
            compact = compact.replace(",", "." if len(tail) in {1, 2} else "")
        elif compact.count(",") > 1:
            compact = compact.replace(",", "")
        try:
            values.append(Decimal(compact))
        except Exception:
            continue
    return values


def contains_number(text: str, expected: Decimal | int, tolerance: Decimal = Decimal("0.01")) -> bool:
    target = Decimal(str(expected))
    return any(abs(value - target) <= tolerance for value in numeric_values(text))


def load_expected(input_dir: Path) -> tuple[list[dict], list[dict], int, Decimal]:
    orders_path = input_dir / ORDERS_NAME
    customers_path = input_dir / CUSTOMERS_NAME
    required = ((orders_path, ORDERS_SHA256), (customers_path, CUSTOMERS_SHA256))
    for path, expected_hash in required:
        if not path.is_file():
            raise FileNotFoundError(f"缺少评测输入：{path}")
        actual_hash = digest(path)
        if actual_hash != expected_hash:
            raise ValueError(f"{path.name} SHA256 不匹配：{actual_hash}")

    with customers_path.open(encoding="utf-8", newline="") as stream:
        customers = list(csv.DictReader(stream))
    with orders_path.open(encoding="utf-8", newline="") as stream:
        orders = list(csv.DictReader(stream))
    if len(customers) != 1451 or len(orders) != 2579:
        raise ValueError(f"输入行数异常：customers={len(customers)}, orders={len(orders)}")

    customer_by_id = {row["CUSTOMER_ID"]: row for row in customers}
    regions: dict[str, list] = defaultdict(lambda: [0, Decimal("0")])
    segments: dict[str, list] = defaultdict(lambda: [0, Decimal("0")])
    for order in orders:
        customer = customer_by_id.get(order["CUSTOMER_ID"])
        if customer is None:
            raise ValueError(f"客户维表缺少 {order['CUSTOMER_ID']}")
        if not ("2024-10-01" <= order["ORDER_DATE"] < "2025-01-01"):
            raise ValueError(f"存在非 Q4 2024 订单：{order['ORDER_ID']}")
        amount = Decimal(order["TOTAL_AMOUNT"])
        for accumulator, key in ((regions, customer["REGION"]), (segments, customer["SEGMENT"])):
            accumulator[key][0] += 1
            accumulator[key][1] += amount

    regional_rows = []
    for name in sorted(regions):
        orders_count, raw_revenue = regions[name]
        revenue = quantize(raw_revenue, "0.01")
        regional_rows.append(
            {
                "name": name,
                "orders": orders_count,
                "revenue": revenue,
                "avg": quantize(revenue / orders_count, "0.01"),
            }
        )
    total_orders = sum(row["orders"] for row in regional_rows)
    total_revenue = quantize(sum((row["revenue"] for row in regional_rows), Decimal("0")), "0.01")

    segment_rows = []
    for name in sorted(segments):
        orders_count, raw_revenue = segments[name]
        revenue = quantize(raw_revenue, "0.01")
        segment_rows.append(
            {
                "name": name,
                "orders": orders_count,
                "revenue": revenue,
                "share": quantize(revenue / total_revenue * 100, "0.1"),
            }
        )
    if len(regional_rows) != 5 or len(segment_rows) != 4 or total_orders != 2579 or total_revenue != Decimal("411857.18"):
        raise ValueError("从静态输入重算出的上游真值异常")
    return regional_rows, segment_rows, total_orders, total_revenue


def run_libreoffice(source: Path, target_ext: str, output_dir: Path) -> tuple[bool, str, Path]:
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    target = output_dir / f"{source.stem}.{target_ext}"
    if executable is None:
        return False, "缺少 LibreOffice", target
    output_dir.mkdir(parents=True, exist_ok=True)
    profile = output_dir / "lo-profile"
    completed = subprocess.run(
        [
            executable,
            "--headless",
            f"-env:UserInstallation={profile.resolve().as_uri()}",
            "--convert-to",
            target_ext,
            "--outdir",
            str(output_dir),
            str(source),
        ],
        capture_output=True,
        text=True,
        timeout=180,
        check=False,
    )
    ok = completed.returncode == 0 and target.is_file()
    evidence = (completed.stderr or completed.stdout or "LibreOffice 无输出")[-800:]
    return ok, evidence, target


def workbook_views(path: Path):
    raw = load_workbook(path, data_only=False)
    cached = load_workbook(path, data_only=True)
    recalc_ok = False
    recalc_message = "未重算"
    temp = tempfile.TemporaryDirectory(prefix="cowork-evo-q4-xlsx-")
    output_dir = Path(temp.name) / "out"
    ok, message, recalculated_path = run_libreoffice(path, "xlsx", output_dir)
    if ok:
        cached.close()
        cached = load_workbook(recalculated_path, data_only=True)
        recalc_ok = True
        recalc_message = "LibreOffice 重算及另存成功"
    else:
        recalc_message = message
    return raw, cached, temp, recalc_ok, recalc_message


def actual_numeric(raw_cell, cached_cell) -> Decimal | None:
    for value in (cached_cell.value, raw_cell.value):
        if isinstance(value, bool) or value is None:
            continue
        if isinstance(value, (int, float, Decimal)):
            return Decimal(str(value))
    return None


def populated_rows(ws) -> list[list]:
    rows = [list(row) for row in ws.iter_rows()]
    while rows and all(cell.value is None for cell in rows[-1]):
        rows.pop()
    return rows


def evaluate_regional_sheet(raw_ws, cached_ws, expected: list[dict]) -> tuple[bool, str, dict[str, tuple[Decimal, Decimal, Decimal]]]:
    rows = populated_rows(raw_ws)
    if not rows:
        return False, "Regional Summary 为空", {}
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in rows[0]]
    if headers != REGIONAL_HEADERS:
        return False, f"表头={headers}", {}
    if len(rows) != 6 or raw_ws.max_column != 4:
        return False, f"数据区尺寸={len(rows)}x{raw_ws.max_column}，期望=6x4", {}
    observed = {}
    errors = []
    for index, truth in enumerate(expected, start=2):
        raw_row = raw_ws[index]
        cached_row = cached_ws[index]
        name = str(raw_row[0].value or "").strip()
        orders = actual_numeric(raw_row[1], cached_row[1])
        revenue = actual_numeric(raw_row[2], cached_row[2])
        avg = actual_numeric(raw_row[3], cached_row[3])
        if name != truth["name"]:
            errors.append(f"{index}行地区={name!r}，期望={truth['name']!r}")
            continue
        if orders is None or orders != truth["orders"]:
            errors.append(f"{name} orders={orders}，期望={truth['orders']}")
        if revenue is None or abs(revenue - truth["revenue"]) > Decimal("0.01"):
            errors.append(f"{name} revenue={revenue}，期望={truth['revenue']}")
        if avg is None or abs(avg - truth["avg"]) > Decimal("0.01"):
            errors.append(f"{name} avg={avg}，期望={truth['avg']}")
        if orders is not None and revenue is not None and avg is not None:
            observed[name] = (orders, revenue, avg)
    return not errors, "5 个地区逐行重算相符" if not errors else "；".join(errors), observed


def display_share(raw_cell, cached_cell) -> Decimal | None:
    value = actual_numeric(raw_cell, cached_cell)
    if value is None:
        return None
    number_format = str(raw_cell.number_format or "")
    if "%" in number_format and abs(value) <= 1:
        return value * 100
    return value


def evaluate_segment_sheet(raw_ws, cached_ws, expected: list[dict]) -> tuple[bool, str, dict[str, tuple[Decimal, Decimal, Decimal]]]:
    rows = populated_rows(raw_ws)
    if not rows:
        return False, "Segment Analysis 为空", {}
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in rows[0]]
    if headers != SEGMENT_HEADERS:
        return False, f"表头={headers}", {}
    if len(rows) != 5 or raw_ws.max_column != 4:
        return False, f"数据区尺寸={len(rows)}x{raw_ws.max_column}，期望=5x4", {}
    observed = {}
    errors = []
    for index, truth in enumerate(expected, start=2):
        raw_row = raw_ws[index]
        cached_row = cached_ws[index]
        name = str(raw_row[0].value or "").strip()
        orders = actual_numeric(raw_row[1], cached_row[1])
        revenue = actual_numeric(raw_row[2], cached_row[2])
        share = display_share(raw_row[3], cached_row[3])
        if name != truth["name"]:
            errors.append(f"{index}行分群={name!r}，期望={truth['name']!r}")
            continue
        if orders is None or orders != truth["orders"]:
            errors.append(f"{name} orders={orders}，期望={truth['orders']}")
        if revenue is None or abs(revenue - truth["revenue"]) > Decimal("0.01"):
            errors.append(f"{name} revenue={revenue}，期望={truth['revenue']}")
        if share is None or abs(share - truth["share"]) > Decimal("0.1"):
            errors.append(f"{name} share={share}，期望={truth['share']}")
        if orders is not None and revenue is not None and share is not None:
            observed[name] = (orders, revenue, share)
    return not errors, "4 个分群逐行重算相符" if not errors else "；".join(errors), observed


def excel_style_check(raw_wb) -> tuple[bool, str]:
    checks = []
    for sheet_name in ("Regional Summary", "Segment Analysis"):
        ws = raw_wb[sheet_name]
        header = list(ws[1])[:4]
        header_bold = all(bool(cell.font.bold) for cell in header)
        header_fill = all(cell.fill.fill_type == "solid" for cell in header)
        filter_or_table = bool(ws.auto_filter.ref) or bool(ws.tables)
        freeze = str(ws.freeze_panes) == "A2"
        width_ok = (ws.column_dimensions["A"].width or 0) >= 20
        revenue_format = str(ws["C2"].number_format or "")
        numeric_format = "0" in revenue_format and revenue_format != "General"
        checks.append((sheet_name, header_bold, header_fill, filter_or_table, freeze, width_ok, numeric_format))
    ok = all(all(row[1:]) for row in checks)
    return ok, f"样式检查={checks}"


def slide_text(slide) -> str:
    return "\n".join(
        shape.text
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    )


def top_titles(slide, slide_height: int, max_top_ratio: float = 0.2) -> list[str]:
    return [
        normalize_text(shape.text)
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
        and shape.text.strip()
        and shape.top <= slide_height * max_top_ratio
    ]


def ppt_structure_check(prs, expected_regions: list[dict]) -> tuple[bool, str]:
    if len(prs.slides) != 7:
        return False, f"页数={len(prs.slides)}"
    title_text = normalize_text(slide_text(prs.slides[0]))
    if "q4 2024 sales executive briefing" not in title_text:
        return False, "标题页缺少指定标题"
    errors = []
    for index, truth in enumerate(expected_regions, start=1):
        title_candidates = top_titles(prs.slides[index], int(prs.slide_height), 0.2)
        if normalize_text(truth["name"]) not in title_candidates:
            errors.append(f"第 {index + 1} 页顶部标题不是 {truth['name']}")
    final_text = normalize_text(slide_text(prs.slides[-1]))
    if "q4 2024" not in final_text or "summary" not in final_text:
        errors.append("最后一页缺少 Q4 2024 Summary")
    return not errors, "标题页、5 个地区页和汇总页顺序正确" if not errors else "；".join(errors)


def ppt_region_metrics_check(prs, expected_regions: list[dict]) -> tuple[bool, str]:
    if len(prs.slides) < 6:
        return False, f"页数不足，无法检查 5 个地区页：{len(prs.slides)}"
    errors = []
    for index, truth in enumerate(expected_regions, start=1):
        text = slide_text(prs.slides[index])
        if normalize_text(truth["name"]) not in normalize_text(text):
            errors.append(f"{truth['name']} 页缺少地区名")
        if not contains_number(text, truth["orders"], Decimal("0")):
            errors.append(f"{truth['name']} 页缺少 orders={truth['orders']}")
        if not contains_number(text, truth["revenue"], Decimal("0.01")):
            errors.append(f"{truth['name']} 页缺少 revenue={truth['revenue']}")
    return not errors, "5 个地区页的订单和收入均正确" if not errors else "；".join(errors)


def ppt_summary_check(prs, total_orders: int, total_revenue: Decimal) -> tuple[bool, str]:
    if not prs.slides:
        return False, "PPTX 没有页面"
    text = slide_text(prs.slides[-1])
    orders_ok = contains_number(text, total_orders, Decimal("0"))
    revenue_ok = contains_number(text, total_revenue, Decimal("0.01"))
    return orders_ok and revenue_ok, f"总订单={orders_ok}，总收入={revenue_ok}"


def ppt_editability_check(prs) -> tuple[bool, str, bool]:
    slide_width = int(prs.slide_width)
    slide_height = int(prs.slide_height)
    ratio_ok = abs(slide_width / slide_height - 16 / 9) <= 0.01
    native_shapes = 0
    text_shapes = 0
    large_pictures = 0
    out_of_bounds_text = []
    for slide_index, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                if shape.width * shape.height >= slide_width * slide_height * 0.60:
                    large_pictures += 1
            else:
                native_shapes += 1
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                text_shapes += 1
                tolerance = 36000
                if (
                    shape.left < -tolerance
                    or shape.top < -tolerance
                    or shape.left + shape.width > slide_width + tolerance
                    or shape.top + shape.height > slide_height + tolerance
                ):
                    out_of_bounds_text.append(f"{slide_index}:{shape.name}")
    flattened = large_pictures > 0 or native_shapes < 35 or text_shapes < 25
    ok = ratio_ok and not flattened and not out_of_bounds_text
    return (
        ok,
        f"16:9={ratio_ok}，原生对象={native_shapes}，文字对象={text_shapes}，"
        f"整页大图={large_pictures}，越界文字={out_of_bounds_text}",
        flattened,
    )


def pdf_pages(path: Path) -> tuple[list[str], str]:
    pdftotext = shutil.which("pdftotext")
    pdfinfo = shutil.which("pdfinfo")
    if not pdftotext or not pdfinfo:
        raise RuntimeError("缺少 pdftotext 或 pdfinfo")
    info = subprocess.run([pdfinfo, str(path)], capture_output=True, text=True, timeout=60, check=False)
    if info.returncode != 0:
        raise RuntimeError((info.stderr or "pdfinfo 失败")[-800:])
    text = subprocess.run(
        [pdftotext, "-layout", str(path), "-"],
        capture_output=True,
        text=True,
        timeout=60,
        check=False,
    )
    if text.returncode != 0:
        raise RuntimeError((text.stderr or "pdftotext 失败")[-800:])
    pages = [page for page in text.stdout.split("\f") if page.strip()]
    return pages, info.stdout


def page_core_check(pages: list[str], expected_regions: list[dict], total_orders: int, total_revenue: Decimal) -> tuple[bool, str]:
    if len(pages) != 7:
        return False, f"PDF 页数={len(pages)}"
    errors = []
    if "q4 2024 sales executive briefing" not in normalize_text(pages[0]):
        errors.append("首页缺少指定标题")
    for index, truth in enumerate(expected_regions, start=1):
        page = pages[index]
        if normalize_text(truth["name"]) not in normalize_text(page):
            errors.append(f"PDF 第 {index + 1} 页缺少 {truth['name']}")
        if not contains_number(page, truth["orders"], Decimal("0")):
            errors.append(f"PDF {truth['name']} 页缺少订单数")
        if not contains_number(page, truth["revenue"], Decimal("0.01")):
            errors.append(f"PDF {truth['name']} 页缺少收入")
    if not contains_number(pages[-1], total_orders, Decimal("0")):
        errors.append("PDF 汇总页缺少总订单")
    if not contains_number(pages[-1], total_revenue, Decimal("0.01")):
        errors.append("PDF 汇总页缺少总收入")
    return not errors, "PDF 7 页核心内容与真值一致" if not errors else "；".join(errors)


def pdf_page_size_ok(info: str) -> bool:
    match = re.search(r"Page size:\s*([\d.]+)\s+x\s+([\d.]+)\s+pts", info)
    if not match:
        return False
    width, height = map(float, match.groups())
    return abs(width / height - 16 / 9) <= 0.01


def cross_artifact_check(
    regional_observed: dict[str, tuple[Decimal, Decimal, Decimal]],
    expected_regions: list[dict],
    prs,
    pages: list[str],
    total_orders: int,
    total_revenue: Decimal,
) -> tuple[bool, str]:
    errors = []
    for index, truth in enumerate(expected_regions, start=1):
        observed = regional_observed.get(truth["name"])
        if observed is None:
            errors.append(f"Excel 缺少 {truth['name']}")
            continue
        orders, revenue, _ = observed
        ppt_text = slide_text(prs.slides[index]) if index < len(prs.slides) else ""
        pdf_text = pages[index] if index < len(pages) else ""
        if not (contains_number(ppt_text, orders, Decimal("0")) and contains_number(pdf_text, orders, Decimal("0"))):
            errors.append(f"{truth['name']} 订单未跨文件一致")
        if not (contains_number(ppt_text, revenue, Decimal("0.01")) and contains_number(pdf_text, revenue, Decimal("0.01"))):
            errors.append(f"{truth['name']} 收入未跨文件一致")
    if len(prs.slides) >= 1 and pages:
        ppt_summary = slide_text(prs.slides[-1])
        pdf_summary = pages[-1]
        if not all(contains_number(text, total_orders, Decimal("0")) for text in (ppt_summary, pdf_summary)):
            errors.append("总订单未跨文件一致")
        if not all(contains_number(text, total_revenue, Decimal("0.01")) for text in (ppt_summary, pdf_summary)):
            errors.append("总收入未跨文件一致")
    return not errors, "Excel、PPTX、PDF 的 5 个地区和季度总计均一致" if not errors else "；".join(errors)


def criterion(identifier: str, triggered_or_passed: bool, evidence: str) -> dict:
    spec = CRITERION_SPECS[identifier]
    return {
        "criterion_id": identifier,
        "description": spec["description"],
        "weight": spec["weight"],
        "type": spec["type"],
        "passed": bool(triggered_or_passed),
        "evidence": evidence,
        "method": "deterministic",
    }


def build_result(results: list[dict], artifact_dir: Path) -> dict:
    present = {item["criterion_id"] for item in results}
    for identifier in CRITERION_SPECS:
        if identifier not in present:
            spec = CRITERION_SPECS[identifier]
            results.append(
                criterion(
                    identifier,
                    spec["type"] == "penalty",
                    "因交付物缺失、损坏或前置门槛失败，本判据未能执行",
                )
            )
    results.sort(key=lambda item: item["criterion_id"])

    positive_total = sum(item["weight"] for item in results if item["weight"] > 0)
    earned = sum(item["weight"] for item in results if item["weight"] > 0 and item["passed"])
    penalties = sum(item["weight"] for item in results if item["weight"] < 0 and item["passed"])
    reward = max(0.0, min(1.0, (earned + penalties) / positive_total))
    hurdle_failures = []
    for item in results:
        if item["type"] == "hurdle" and not item["passed"]:
            hurdle_failures.append(item["criterion_id"])
            reward = min(reward, float(CRITERION_SPECS[item["criterion_id"]]["cap"]))

    artifact_hashes = []
    for filename in (XLSX_NAME, PPTX_NAME, PDF_NAME):
        path = artifact_dir / filename
        if path.is_file():
            value = digest(path)
            if value not in artifact_hashes:
                artifact_hashes.append(value)
    if not artifact_hashes:
        artifact_hashes = [hashlib.sha256(b"").hexdigest()]

    normalized = []
    for item in results:
        score = (0.0 if item["passed"] else 1.0) if item["type"] == "penalty" else (1.0 if item["passed"] else 0.0)
        normalized.append(
            {
                "criterion_id": item["criterion_id"],
                "score": score,
                "evidence": item["evidence"],
                "method": item["method"],
                "raw": {
                    "description": item["description"],
                    "weight": item["weight"],
                    "type": item["type"],
                    "passed": item["passed"],
                },
            }
        )
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": artifact_hashes,
        "criterion_results": normalized,
        "reward": round(reward, 6),
        "passed": reward >= 0.8,
        "hurdle_failures": hurdle_failures,
    }


def evaluate(input_dir: Path, artifact_dir: Path) -> dict:
    results: list[dict] = []
    expected_regions, expected_segments, total_orders, total_revenue = load_expected(input_dir)
    xlsx_path = artifact_dir / XLSX_NAME
    pptx_path = artifact_dir / PPTX_NAME
    pdf_path = artifact_dir / PDF_NAME

    missing = [path.name for path in (xlsx_path, pptx_path, pdf_path) if not path.is_file()]
    if missing:
        results.append(criterion("R001", False, f"缺少文件={missing}"))
        results.append(criterion("R013", True, f"缺少文件={missing}"))
        return build_result(results, artifact_dir)

    raw_wb = cached_wb = None
    temp_recalc = None
    try:
        raw_wb, cached_wb, temp_recalc, recalc_ok, recalc_message = workbook_views(xlsx_path)
        prs = Presentation(pptx_path)
        pages, pdf_info = pdf_pages(pdf_path)
    except Exception as exc:
        if raw_wb:
            raw_wb.close()
        if cached_wb:
            cached_wb.close()
        if temp_recalc:
            temp_recalc.cleanup()
        results.append(criterion("R001", False, f"交付物无法打开：{exc}"))
        results.append(criterion("R013", True, f"交付物无法打开：{exc}"))
        return build_result(results, artifact_dir)

    file_ok = (
        raw_wb.sheetnames == ["Regional Summary", "Segment Analysis"]
        and len(prs.slides) == 7
        and len(pages) == 7
        and pdf_path.stat().st_size >= 10_000
    )
    results.append(
        criterion(
            "R001",
            file_ok,
            f"Excel 表={raw_wb.sheetnames}，PPTX 页={len(prs.slides)}，PDF 页={len(pages)}，PDF={pdf_path.stat().st_size} bytes",
        )
    )

    regional_ok, regional_evidence, regional_observed = evaluate_regional_sheet(
        raw_wb["Regional Summary"] if "Regional Summary" in raw_wb.sheetnames else raw_wb.active,
        cached_wb["Regional Summary"] if "Regional Summary" in cached_wb.sheetnames else cached_wb.active,
        expected_regions,
    )
    results.append(criterion("R002", regional_ok, regional_evidence))

    if "Segment Analysis" in raw_wb.sheetnames and "Segment Analysis" in cached_wb.sheetnames:
        segment_ok, segment_evidence, _ = evaluate_segment_sheet(
            raw_wb["Segment Analysis"], cached_wb["Segment Analysis"], expected_segments
        )
    else:
        segment_ok, segment_evidence = False, "缺少 Segment Analysis"
    results.append(criterion("R003", segment_ok, segment_evidence))

    if all(name in raw_wb.sheetnames for name in ("Regional Summary", "Segment Analysis")):
        style_ok, style_evidence = excel_style_check(raw_wb)
    else:
        style_ok, style_evidence = False, "工作表结构不完整"
    results.append(criterion("R004", style_ok, style_evidence))

    structure_ok, structure_evidence = ppt_structure_check(prs, expected_regions)
    results.append(criterion("R005", structure_ok, structure_evidence))
    region_ppt_ok, region_ppt_evidence = ppt_region_metrics_check(prs, expected_regions)
    results.append(criterion("R006", region_ppt_ok, region_ppt_evidence))
    summary_ok, summary_evidence = ppt_summary_check(prs, total_orders, total_revenue)
    results.append(criterion("R007", summary_ok, summary_evidence))
    editable_ok, editable_evidence, flattened = ppt_editability_check(prs)
    results.append(criterion("R008", editable_ok, editable_evidence))

    pdf_core_ok, pdf_core_evidence = page_core_check(pages, expected_regions, total_orders, total_revenue)
    pdf_ok = pdf_core_ok and pdf_page_size_ok(pdf_info)
    results.append(criterion("R009", pdf_ok, f"{pdf_core_evidence}；16:9={pdf_page_size_ok(pdf_info)}"))

    cross_ok, cross_evidence = cross_artifact_check(
        regional_observed, expected_regions, prs, pages, total_orders, total_revenue
    )
    results.append(criterion("R010", cross_ok, cross_evidence))

    with tempfile.TemporaryDirectory(prefix="cowork-evo-q4-ppt-") as temp_dir:
        rendered_ok, render_message, rendered_pdf = run_libreoffice(pptx_path, "pdf", Path(temp_dir))
        if rendered_ok:
            try:
                rendered_pages, _ = pdf_pages(rendered_pdf)
                rendered_core_ok, rendered_core_evidence = page_core_check(
                    rendered_pages, expected_regions, total_orders, total_revenue
                )
            except Exception as exc:
                rendered_core_ok, rendered_core_evidence = False, str(exc)
        else:
            rendered_core_ok, rendered_core_evidence = False, render_message
    results.append(criterion("R011", rendered_ok and rendered_core_ok, rendered_core_evidence))
    results.append(criterion("R012", recalc_ok, recalc_message))

    broken_or_empty = not file_ok
    results.append(criterion("R013", broken_or_empty, f"文件门槛失败={broken_or_empty}"))
    pdf_ppt_mismatch = not (pdf_core_ok and rendered_core_ok and cross_ok)
    results.append(
        criterion(
            "R014",
            flattened or pdf_ppt_mismatch,
            f"压平={flattened}，PDF/PPTX 核心内容不一致={pdf_ppt_mismatch}",
        )
    )

    raw_wb.close()
    cached_wb.close()
    temp_recalc.cleanup()
    return build_result(results, artifact_dir)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path)
    args = parser.parse_args()
    result = evaluate(args.input_dir, args.artifact_dir)
    rendered = json.dumps(result, ensure_ascii=False, indent=2)
    print(rendered)
    if args.result:
        args.result.parent.mkdir(parents=True, exist_ok=True)
        args.result.write_text(rendered + "\n", encoding="utf-8")
    raise SystemExit(0 if result["passed"] else 1)


if __name__ == "__main__":
    main()
