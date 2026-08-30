#!/usr/bin/env python3
"""逐页制作 OfficeBench 3-2 NVIDIA Q1 FY2024 财务简报 Golden。"""

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


SOURCE_SHA256 = "33b606d43e02d2808822258ce6b4edfcd7c7bb7ed471f5395583286707a45626"

BLACK = RGBColor(0x10, 0x17, 0x14)
INK = RGBColor(0x1E, 0x29, 0x25)
GREEN = RGBColor(0x76, 0xB9, 0x00)
DARK_GREEN = RGBColor(0x37, 0x62, 0x16)
LIME = RGBColor(0xC8, 0xF0, 0x72)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SURFACE = RGBColor(0xF5, 0xF7, 0xF5)
LINE = RGBColor(0xD9, 0xE0, 0xDB)
MUTED = RGBColor(0x65, 0x70, 0x6A)
RED = RGBColor(0xC4, 0x32, 0x2F)
AMBER = RGBColor(0xB8, 0x72, 0x00)
BLUE = RGBColor(0x2C, 0x68, 0xA0)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def add_text(slide, text, left, top, width, height, *, size=15, color=INK,
             bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = str(text)
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    return shape


def add_rect(slide, left, top, width, height, fill, *, line=None, rounded=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def background(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def header(slide, title, kicker, page):
    add_text(slide, kicker.upper(), 0.65, 0.28, 7.5, 0.24, size=9, color=DARK_GREEN, bold=True)
    add_text(slide, title, 0.65, 0.57, 11.7, 0.55, size=27, color=BLACK, bold=True)
    add_text(slide, f"{page:02d} / 06", 11.65, 0.33, 0.95, 0.2, size=9, color=MUTED, align=PP_ALIGN.RIGHT)
    add_rect(slide, 0.65, 1.18, 12.03, 0.02, LINE, rounded=False)


def footer(slide, page, pages, dark=False):
    color = RGBColor(0xA8, 0xB2, 0xAC) if dark else MUTED
    add_text(slide, f"Source: OfficeBench financial_report.pdf, pp. {pages}  ·  NVIDIA Q1 FY2024 release, 24 May 2023", 0.65, 7.13, 10.8, 0.18, size=7.5, color=color)
    add_text(slide, str(page), 12.15, 7.13, 0.45, 0.18, size=8, color=color, align=PP_ALIGN.RIGHT)


def metric(slide, value, label, detail, left, top, width, accent=GREEN):
    add_rect(slide, left, top, width, 1.17, WHITE, line=LINE)
    add_rect(slide, left, top, 0.07, 1.17, accent, rounded=False)
    add_text(slide, value, left + 0.22, top + 0.12, width - 0.3, 0.38, size=22, color=BLACK, bold=True)
    add_text(slide, label.upper(), left + 0.22, top + 0.57, width - 0.3, 0.2, size=8.5, color=MUTED, bold=True)
    add_text(slide, detail, left + 0.22, top + 0.84, width - 0.3, 0.2, size=8.5, color=accent, bold=True)


def style_table(table, widths, *, font_size=10):
    for column, width in zip(table.columns, widths, strict=True):
        column.width = Inches(width)
    for column, cell in enumerate(table.rows[0].cells):
        cell.fill.solid(); cell.fill.fore_color.rgb = BLACK
        for run in cell.text_frame.paragraphs[0].runs:
            run.font.name = "Aptos"; run.font.size = Pt(font_size); run.font.bold = True; run.font.color.rgb = WHITE
    for row_index in range(1, len(table.rows)):
        for cell in table.rows[row_index].cells:
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE if row_index % 2 else SURFACE
            cell.margin_left = cell.margin_right = Inches(0.07)
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Aptos"; run.font.size = Pt(font_size); run.font.color.rgb = INK


def build(input_dir: Path, output: Path) -> None:
    source = input_dir / "financial_report.pdf"
    if not source.is_file() or sha256(source) != SOURCE_SHA256:
        raise ValueError("OfficeBench 3-2 冻结 PDF 缺失或 SHA256 不匹配")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1. Executive cover
    slide = prs.slides.add_slide(blank); background(slide, BLACK)
    add_text(slide, "NVIDIA  /  Q1 FY2024", 0.72, 0.62, 5.5, 0.26, size=10, color=LIME, bold=True)
    add_text(slide, "Sequential acceleration,\nwith a mixed year-on-year base", 0.72, 1.18, 8.0, 1.32, size=34, color=WHITE, bold=True)
    add_text(slide, "Management financial briefing  ·  Quarter ended 30 April 2023", 0.75, 2.78, 7.3, 0.34, size=14, color=RGBColor(0xC8, 0xD1, 0xCB))
    add_rect(slide, 9.22, 0.95, 3.35, 2.52, RGBColor(0x20, 0x2B, 0x25), line=RGBColor(0x3B, 0x4A, 0x42))
    add_text(slide, "$7.19B", 9.48, 1.36, 2.85, 0.58, size=31, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Q1 REVENUE", 9.48, 2.03, 2.85, 0.22, size=9, color=RGBColor(0xB8, 0xC4, 0xBC), bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "+19% Q/Q  ·  −13% Y/Y", 9.48, 2.55, 2.85, 0.28, size=13, color=LIME, bold=True, align=PP_ALIGN.CENTER)
    metric(slide, "$4.28B", "Data Center revenue", "record  ·  +18% Q/Q", 0.72, 4.73, 3.75, GREEN)
    metric(slide, "$0.82", "GAAP diluted EPS", "+44% Q/Q  ·  +28% Y/Y", 4.79, 4.73, 3.75, GREEN)
    metric(slide, "$11.00B", "Q2 revenue outlook", "±2%", 8.86, 4.73, 3.75, GREEN)
    add_text(slide, "FACT-BASED SUMMARY  ·  NO INVESTMENT RECOMMENDATION", 0.75, 6.37, 7.2, 0.24, size=9, color=RGBColor(0xA8, 0xB2, 0xAC), bold=True)
    footer(slide, 1, "1–2", dark=True)

    # 2. Headline scorecard and native revenue chart
    slide = prs.slides.add_slide(blank); background(slide, WHITE); header(slide, "Headline scorecard", "Quarter at a glance", 2)
    data = ChartData(); data.categories = ["Q1 FY23", "Q4 FY23", "Q1 FY24"]; data.add_series("Revenue ($M)", [8288, 6051, 7192])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.72), Inches(1.78), Inches(6.05), Inches(4.56), data).chart
    chart.has_legend = False; chart.has_title = False; chart.value_axis.minimum_scale = 0; chart.value_axis.maximum_scale = 9000
    plot = chart.plots[0]; plot.has_data_labels = True; plot.data_labels.show_value = True; plot.data_labels.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
    plot.series[0].format.fill.solid(); plot.series[0].format.fill.fore_color.rgb = GREEN
    add_text(slide, "Revenue reset, then rebound", 0.78, 1.36, 4.2, 0.28, size=14, color=BLACK, bold=True)
    cards = [
        ("64.6%", "GAAP gross margin", "+1.3 pts Q/Q  ·  −0.9 pts Y/Y", GREEN),
        ("$2.14B", "GAAP operating income", "+70% Q/Q  ·  +15% Y/Y", GREEN),
        ("$2.04B", "GAAP net income", "+44% Q/Q  ·  +26% Y/Y", GREEN),
        ("$2.51B", "GAAP operating expenses", "−3% Q/Q  ·  −30% Y/Y", BLUE),
    ]
    for index, (value, label, detail, accent) in enumerate(cards):
        metric(slide, value, label, detail, 7.15, 1.52 + index * 1.22, 5.2, accent)
    add_text(slide, "Derived reading: sequential earnings expanded faster than revenue; revenue remained below Q1 FY2023.", 0.86, 6.48, 11.5, 0.34, size=10.5, color=MUTED)
    footer(slide, 2, "1–2")

    # 3. Exact GAAP table
    slide = prs.slides.add_slide(blank); background(slide, WHITE); header(slide, "GAAP operating leverage", "Reported scorecard", 3)
    rows = [
        ["Metric", "Q1 FY24", "Q4 FY23", "Q1 FY23", "Q/Q", "Y/Y"],
        ["Revenue ($M)", "7,192", "6,051", "8,288", "Up 19%", "Down 13%"],
        ["Gross margin", "64.6%", "63.3%", "65.5%", "Up 1.3 pts", "Down 0.9 pts"],
        ["Operating expenses ($M)", "2,508", "2,576", "3,563", "Down 3%", "Down 30%"],
        ["Operating income ($M)", "2,140", "1,257", "1,868", "Up 70%", "Up 15%"],
        ["Net income ($M)", "2,043", "1,414", "1,618", "Up 44%", "Up 26%"],
        ["Diluted EPS", "$0.82", "$0.57", "$0.64", "Up 44%", "Up 28%"],
    ]
    table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(0.72), Inches(1.55), Inches(11.9), Inches(4.65)).table
    for r, values in enumerate(rows):
        for c, value in enumerate(values): table.cell(r, c).text = value
    style_table(table, [2.9, 1.55, 1.55, 1.55, 1.65, 1.65], font_size=10.5)
    for r in range(1, len(rows)):
        for c in range(1, len(rows[0])):
            table.cell(r, c).text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    for r in range(len(rows)): table.rows[r].height = Inches(0.62 if r else 0.5)
    add_rect(slide, 0.72, 6.38, 11.9, 0.48, SURFACE, line=LINE)
    add_text(slide, "All values above are transcribed from the GAAP Q1 Fiscal 2024 Summary; $ values are in millions except EPS.", 0.92, 6.51, 11.5, 0.2, size=9.5, color=MUTED)
    footer(slide, 3, "1–2")

    # 4. Segment momentum
    slide = prs.slides.add_slide(blank); background(slide, SURFACE); header(slide, "Data Center led the mix", "Platform revenue", 4)
    data = ChartData(); data.categories = ["Data Center", "Gaming", "Professional Visualization", "Automotive"]; data.add_series("Revenue ($B)", [4.28, 2.24, 0.295, 0.296])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.72), Inches(1.55), Inches(6.1), Inches(4.95), data).chart
    chart.has_legend = False; chart.has_title = False; chart.value_axis.minimum_scale = 0; chart.value_axis.maximum_scale = 5
    plot = chart.plots[0]; plot.has_data_labels = True; plot.data_labels.show_value = True; plot.data_labels.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
    plot.series[0].format.fill.solid(); plot.series[0].format.fill.fore_color.rgb = GREEN
    segment_cards = [
        ("DATA CENTER", "$4.28B", "+18% Q/Q  ·  +14% Y/Y", "Record quarter", GREEN),
        ("GAMING", "$2.24B", "+22% Q/Q  ·  −38% Y/Y", "Sequential recovery", AMBER),
        ("PRO VIS", "$295M", "+31% Q/Q  ·  −53% Y/Y", "Small, mixed base", RED),
        ("AUTOMOTIVE", "$296M", "+1% Q/Q  ·  +114% Y/Y", "Record quarter", BLUE),
    ]
    for index, (name, value, change, note, accent) in enumerate(segment_cards):
        top = 1.55 + index * 1.18
        add_rect(slide, 7.12, top, 5.25, 0.98, WHITE, line=LINE)
        add_rect(slide, 7.12, top, 0.07, 0.98, accent, rounded=False)
        add_text(slide, name, 7.38, top + 0.13, 1.65, 0.22, size=9, color=MUTED, bold=True)
        add_text(slide, value, 9.02, top + 0.11, 1.15, 0.32, size=18, color=BLACK, bold=True, align=PP_ALIGN.RIGHT)
        add_text(slide, change, 10.34, top + 0.14, 1.7, 0.24, size=9, color=accent, bold=True, align=PP_ALIGN.RIGHT)
        add_text(slide, note, 7.38, top + 0.53, 4.63, 0.22, size=9.5, color=INK)
    footer(slide, 4, "2")

    # 5. GAAP/non-GAAP and cash conversion
    slide = prs.slides.add_slide(blank); background(slide, WHITE); header(slide, "Earnings quality and cash conversion", "Reconciliation", 5)
    rows = [
        ["Q1 FY2024", "GAAP", "Non-GAAP"],
        ["Gross margin", "64.6%", "66.8%"],
        ["Operating expenses ($M)", "2,508", "1,750"],
        ["Operating income ($M)", "2,140", "3,052"],
        ["Net income ($M)", "2,043", "2,713"],
        ["Diluted EPS", "$0.82", "$1.09"],
    ]
    table = slide.shapes.add_table(len(rows), 3, Inches(0.72), Inches(1.55), Inches(6.45), Inches(4.42)).table
    for r, values in enumerate(rows):
        for c, value in enumerate(values): table.cell(r, c).text = value
    style_table(table, [3.2, 1.6, 1.65], font_size=11)
    for r in range(1, len(rows)):
        table.cell(r, 1).text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
        table.cell(r, 2).text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    for r in range(len(rows)): table.rows[r].height = Inches(0.66 if r else 0.5)
    metric(slide, "$2.911B", "Operating cash flow", "vs $1.731B Q1 FY23", 7.55, 1.55, 4.8, GREEN)
    metric(slide, "$2.643B", "Free cash flow", "vs $1.348B Q1 FY23", 7.55, 2.92, 4.8, GREEN)
    metric(slide, "$99M", "Cash dividends", "returned in Q1 FY24", 7.55, 4.29, 4.8, BLUE)
    add_rect(slide, 0.72, 6.22, 11.63, 0.62, SURFACE, line=LINE)
    add_text(slide, "Non-GAAP measures exclude specified acquisition-related costs, stock compensation and other adjustments; they do not replace GAAP.", 0.94, 6.39, 11.18, 0.28, size=9.5, color=MUTED)
    footer(slide, 5, "3, 5–6")

    # 6. Outlook and watchpoints
    slide = prs.slides.add_slide(blank); background(slide, BLACK)
    add_text(slide, "Q2 FY2024 OUTLOOK  /  FORWARD-LOOKING", 0.72, 0.58, 7.5, 0.25, size=10, color=LIME, bold=True)
    add_text(slide, "$11.00B revenue guide\nraises the near-term run-rate", 0.72, 1.03, 7.8, 1.16, size=31, color=WHITE, bold=True)
    outlook = [
        ("Revenue", "$11.00B", "±2%"),
        ("GAAP gross margin", "68.6%", "±50 bps"),
        ("Non-GAAP gross margin", "70.0%", "±50 bps"),
        ("GAAP / non-GAAP OpEx", "$2.71B / $1.90B", "approx."),
        ("Other income", "$90M", "approx."),
        ("GAAP / non-GAAP tax rate", "14.0%", "±1%"),
    ]
    for index, (label, value, qualifier) in enumerate(outlook):
        row, column = divmod(index, 2); left = 0.72 + column * 4.03; top = 2.72 + row * 1.08
        add_rect(slide, left, top, 3.72, 0.84, RGBColor(0x20, 0x2B, 0x25), line=RGBColor(0x3B, 0x4A, 0x42))
        add_text(slide, label, left + 0.18, top + 0.12, 2.1, 0.22, size=9, color=RGBColor(0xB8, 0xC4, 0xBC), bold=True)
        add_text(slide, value, left + 0.18, top + 0.4, 2.55, 0.28, size=16, color=WHITE, bold=True)
        add_text(slide, qualifier, left + 2.56, top + 0.45, 0.85, 0.2, size=9, color=LIME, bold=True, align=PP_ALIGN.RIGHT)
    add_rect(slide, 8.94, 1.03, 3.62, 4.93, WHITE, line=WHITE)
    add_text(slide, "MANAGEMENT WATCHPOINTS", 9.28, 1.4, 2.92, 0.24, size=9, color=DARK_GREEN, bold=True)
    add_text(slide, "1", 9.28, 1.93, 0.35, 0.32, size=20, color=GREEN, bold=True)
    add_text(slide, "Supply ramp", 9.72, 1.93, 2.1, 0.28, size=15, color=BLACK, bold=True)
    add_text(slide, "Release says supply is increasing to meet surging Data Center demand.", 9.72, 2.3, 2.25, 0.68, size=10, color=INK)
    add_text(slide, "2", 9.28, 3.13, 0.35, 0.32, size=20, color=GREEN, bold=True)
    add_text(slide, "Margin delivery", 9.72, 3.13, 2.1, 0.28, size=15, color=BLACK, bold=True)
    add_text(slide, "Track the stated ±50 bps ranges for both GAAP and non-GAAP gross margin.", 9.72, 3.5, 2.25, 0.68, size=10, color=INK)
    add_text(slide, "3", 9.28, 4.34, 0.35, 0.32, size=20, color=GREEN, bold=True)
    add_text(slide, "Comparison base", 9.72, 4.34, 2.1, 0.28, size=15, color=BLACK, bold=True)
    add_text(slide, "Separate sequential recovery from year-on-year declines in Gaming and Pro Visualization.", 9.72, 4.71, 2.25, 0.76, size=10, color=INK)
    add_text(slide, "Outlook is not a guarantee of future performance; see the source release for risk factors.", 0.75, 6.47, 7.9, 0.35, size=9.5, color=RGBColor(0xA8, 0xB2, 0xAC))
    footer(slide, 6, "1–3, 7", dark=True)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    if len(Presentation(output).slides) != 6:
        raise ValueError("Golden 保存后页数不是 6")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.input_dir, args.output)
    print(f"{args.output}\nsha256={sha256(args.output)}")


if __name__ == "__main__":
    main()
