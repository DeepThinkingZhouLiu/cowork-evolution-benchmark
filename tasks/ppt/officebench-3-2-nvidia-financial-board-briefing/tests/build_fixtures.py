#!/usr/bin/env python3
"""OfficeBench 3-2 PPT Judge 的逐项校准样例。"""
from __future__ import annotations
import argparse, shutil, zipfile
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

OUTPUT_NAME = "NVIDIA_Q1_FY2024_Management_Briefing.pptx"
C = "http://schemas.openxmlformats.org/drawingml/2006/chart"

def out(root: Path, name: str) -> Path:
    path=root/name/OUTPUT_NAME; path.parent.mkdir(parents=True,exist_ok=True); return path

def mutate(golden: Path, output: Path, fn) -> None:
    prs=Presentation(golden); fn(prs); prs.save(output)

def replace_all(prs: Presentation, old: str, new: str) -> None:
    changed=0
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape,"has_text_frame",False) and old in shape.text:
                shape.text=shape.text.replace(old,new); changed+=1
    if not changed: raise ValueError(old)

def find_table(prs: Presentation, first_header: str):
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_table and shape.table.cell(0,0).text == first_header: return shape.table
    raise ValueError(first_header)

def change_chart(prs: Presentation, index: int, value: str) -> None:
    charts=[shape.chart for slide in prs.slides for shape in slide.shapes if shape.shape_type==MSO_SHAPE_TYPE.CHART]
    cache=charts[index]._chartSpace.find(f".//{{{C}}}numCache")
    cache.findall(f"{{{C}}}pt")[0].find(f"{{{C}}}v").text=value

def main() -> None:
    parser=argparse.ArgumentParser(); parser.add_argument("--golden",type=Path,required=True); parser.add_argument("--outdir",type=Path,required=True); args=parser.parse_args()
    (args.outdir/"empty").mkdir(parents=True); out(args.outdir,"corrupt").write_bytes(b"bad\n")
    with zipfile.ZipFile(args.golden) as reader, zipfile.ZipFile(out(args.outdir,"valid-repackage"),"w",compression=zipfile.ZIP_STORED) as writer:
        for item in reader.infolist(): writer.writestr(item.filename,reader.read(item.filename))
    def alt(prs): prs.slides[0].background.fill.solid(); prs.slides[0].background.fill.fore_color.rgb=RGBColor(0x17,0x24,0x20)
    mutate(args.golden,out(args.outdir,"valid-alt-cover"),alt)
    mutate(args.golden,out(args.outdir,"wrong-revenue-chart"),lambda prs: change_chart(prs,0,"9999"))
    def missing_gaap(prs):
        table=find_table(prs,"Metric")
        for cell in table.rows[2].cells: cell.text=""
    mutate(args.golden,out(args.outdir,"missing-gaap-row"),missing_gaap)
    def wrong_gaap(prs): find_table(prs,"Metric").cell(4,1).text="2,240"
    mutate(args.golden,out(args.outdir,"wrong-gaap-value"),wrong_gaap)
    mutate(args.golden,out(args.outdir,"wrong-segment-chart"),lambda prs: change_chart(prs,1,"5.28"))
    mutate(args.golden,out(args.outdir,"wrong-segment-change"),lambda prs: replace_all(prs,"+14% Y/Y","+15% Y/Y"))
    def wrong_nongaap(prs): find_table(prs,"Q1 FY2024").cell(1,2).text="67.8%"
    mutate(args.golden,out(args.outdir,"wrong-non-gaap"),wrong_nongaap)
    mutate(args.golden,out(args.outdir,"wrong-fcf"),lambda prs: replace_all(prs,"$2.643B","$2.700B"))
    mutate(args.golden,out(args.outdir,"wrong-outlook-revenue"),lambda prs: replace_all(prs,"$11.00B","$12.00B"))
    mutate(args.golden,out(args.outdir,"wrong-margin-range"),lambda prs: replace_all(prs,"±50 bps","±75 bps"))
    mutate(args.golden,out(args.outdir,"missing-citations"),lambda prs: replace_all(prs,"Source: OfficeBench financial_report.pdf","Source omitted"))
    def investment(prs):
        box=prs.slides[0].shapes.add_textbox(Inches(8.8),Inches(6.2),Inches(3.2),Inches(.3)); box.text="BUY NVIDIA — OUTPERFORM RATING"
    mutate(args.golden,out(args.outdir,"investment-advice"),investment)
    image=args.outdir/"flat.png"; Image.new("RGB",(1600,900),"white").save(image)
    def flat(prs):
        slide=prs.slides[2]
        for shape in list(slide.shapes): shape._element.getparent().remove(shape._element)
        slide.shapes.add_picture(str(image),0,0,width=prs.slide_width,height=prs.slide_height)
    mutate(args.golden,out(args.outdir,"flattened-gaap-table"),flat)
    def marker(prs):
        box=prs.slides[-1].shapes.add_textbox(Inches(.1),Inches(.1),Inches(2),Inches(.2)); box.text="solution/Golden"
    mutate(args.golden,out(args.outdir,"hardcoded-marker"),marker)
    extra=args.outdir/"extra-output"; extra.mkdir(); shutil.copy2(args.golden,extra/OUTPUT_NAME); (extra/"debug.txt").write_text("debug\n")

if __name__=="__main__": main()
