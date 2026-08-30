#!/usr/bin/env python3
"""OfficeBench 3-2 NVIDIA 财务简报的任务专属确定性 Judge。"""
from __future__ import annotations

import argparse
import hashlib
import json
import posixpath
import re
import subprocess
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "NVIDIA_Q1_FY2024_Management_Briefing.pptx"
SOURCE_SHA256 = "33b606d43e02d2808822258ce6b4edfcd7c7bb7ed471f5395583286707a45626"
TASK_ID = "cowork-evo/officebench-3-2-nvidia-financial-board-briefing"
RUBRIC_VERSION = "officebench-3-2-ppt-adaptation-v0.1"
JUDGE_VERSION = "officebench-3-2-deterministic-v0.1"
SUSPICIOUS = ["solution/golden", "judge-result", "artifact_dir", "officebench-3-2", "/root/"]

GAAP_ROWS = {
    "Revenue ($M)": ("7,192", "6,051", "8,288", "Up 19%", "Down 13%"),
    "Gross margin": ("64.6%", "63.3%", "65.5%", "Up 1.3 pts", "Down 0.9 pts"),
    "Operating expenses ($M)": ("2,508", "2,576", "3,563", "Down 3%", "Down 30%"),
    "Operating income ($M)": ("2,140", "1,257", "1,868", "Up 70%", "Up 15%"),
    "Net income ($M)": ("2,043", "1,414", "1,618", "Up 44%", "Up 26%"),
    "Diluted EPS": ("$0.82", "$0.57", "$0.64", "Up 44%", "Up 28%"),
}
NON_GAAP_ROWS = {
    "Gross margin": ("64.6%", "66.8%"),
    "Operating expenses ($M)": ("2,508", "1,750"),
    "Operating income ($M)": ("2,140", "3,052"),
    "Net income ($M)": ("2,043", "2,713"),
    "Diluted EPS": ("$0.82", "$1.09"),
}
SEGMENT_TOKENS = [
    "DATA CENTER", "$4.28B", "+18% Q/Q", "+14% Y/Y", "Record quarter",
    "GAMING", "$2.24B", "+22% Q/Q", "−38% Y/Y",
    "PRO VIS", "$295M", "+31% Q/Q", "−53% Y/Y",
    "AUTOMOTIVE", "$296M", "+1% Q/Q", "+114% Y/Y",
]
CASH_TOKENS = ["$2.911B", "Operating cash flow", "$2.643B", "Free cash flow", "$99M", "Cash dividends"]
OUTLOOK_TOKENS = [
    "Q2 FY2024 OUTLOOK", "$11.00B", "±2%", "68.6%", "70.0%", "±50 bps",
    "$2.71B / $1.90B", "$90M", "14.0%", "±1%",
]
SOURCE_TOKENS = [
    "Revenue", "7,192", "6,051", "8,288", "64.6%", "2,508", "2,140", "2,043", "0.82",
    "Data Center", "4.28 billion", "Gaming", "2.24 billion", "Professional Visualization", "295 million",
    "Automotive", "296 million", "2,911", "2,643", "11.00 billion", "68.6%", "70.0%", "2.71 billion", "1.90 billion",
]
SPECS = {
    "R001": ("PPTX 完整、16:9 且 6 页", 6, "hurdle"),
    "R002": ("Revenue 原生图准确", 12, "hurdle"),
    "R003": ("GAAP 三期表准确", 22, "hurdle"),
    "R004": ("四业务收入与变化准确", 14, "hurdle"),
    "R005": ("GAAP/Non-GAAP 与现金流准确", 16, "hurdle"),
    "R006": ("Q2 outlook 完整准确", 14, "hurdle"),
    "R007": ("来源与风险边界完整", 6, "hurdle"),
    "R008": ("原生可编辑且无标记/宏/外链", 8, "hurdle"),
    "R009": ("6 页可渲染且输出干净", 2, "positive"),
    "R010": ("无严重事实或可编辑性错误", -20, "penalty"),
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def norm(value) -> str:
    return re.sub(r"\s+", " ", str(value).replace("\u00a0", " ").replace("\v", " ")).strip()


def locate_source() -> Path | None:
    for path in (Path("/app/input_files/financial_report.pdf"), Path(__file__).resolve().parents[1] / "data/input_files/financial_report.pdf"):
        if path.is_file(): return path
    return None


def source_check() -> tuple[bool, str]:
    source = locate_source()
    if source is None or sha256(source) != SOURCE_SHA256:
        return False, "冻结 PDF 缺失或 SHA256 错误"
    run = subprocess.run(["pdftotext", "-layout", str(source), "-"], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=120)
    text = run.stdout
    missing = [token for token in SOURCE_TOKENS if token not in text]
    return run.returncode == 0 and not missing, f"pdf_sha256={sha256(source)} missing={missing}"


def package_checks(path: Path):
    evidence = {"broken_member": None, "missing_targets": [], "macros": [], "external_targets": []}
    try:
        with zipfile.ZipFile(path) as archive:
            evidence["broken_member"] = archive.testzip(); names = set(archive.namelist())
            evidence["macros"] = [name for name in names if name.lower().endswith("vbaproject.bin")]
            for rel_name in sorted(name for name in names if name.endswith(".rels")):
                root = etree.fromstring(archive.read(rel_name))
                owner = "" if rel_name == "_rels/.rels" else posixpath.join(rel_name.split("/_rels/", 1)[0], rel_name.split("/_rels/", 1)[1][:-5])
                base = posixpath.dirname(owner)
                for relationship in root:
                    target = relationship.attrib.get("Target", "")
                    if relationship.attrib.get("TargetMode") == "External":
                        evidence["external_targets"].append(target); continue
                    resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                    if resolved not in names: evidence["missing_targets"].append(f"{rel_name} -> {resolved}")
        ok = evidence["broken_member"] is None and not evidence["missing_targets"] and not evidence["macros"]
        return ok, evidence
    except Exception as error:
        evidence["exception"] = str(error); return False, evidence


def slide_text(slide) -> str:
    return "\n".join(shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False) and not shape.has_table)


def deck_text(prs: Presentation) -> str:
    return "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False))


def chart_mappings(prs: Presentation):
    result = []
    for slide_number, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.CHART: continue
            try:
                chart = shape.chart; categories = [norm(item.label) for item in chart.plots[0].categories]
                values = [float(value) for value in chart.series[0].values]
                result.append((slide_number, dict(zip(categories, values, strict=True))))
            except Exception as error:
                result.append((slide_number, {"error": str(error)}))
    return result


def table_candidates(prs: Presentation):
    candidates = []
    for slide_number, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if not shape.has_table: continue
            table = shape.table
            rows = [[norm(cell.text) for cell in row.cells] for row in table.rows]
            candidates.append((slide_number, rows))
    return candidates


def find_table(candidates, headers, expected):
    matches = []
    for slide_number, rows in candidates:
        if not rows or rows[0] != headers: continue
        observed = {row[0]: tuple(row[1:]) for row in rows[1:] if row and row[0]}
        if observed == expected: matches.append(slide_number)
    return matches


def render(path: Path, out: Path):
    out.mkdir(parents=True, exist_ok=True); profile = out / "profile"; profile.mkdir(exist_ok=True)
    run = subprocess.run(["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf", "--outdir", str(out), str(path)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300)
    pdf = out / f"{path.stem}.pdf"
    if run.returncode or not pdf.is_file(): return 0, run.stdout[-500:]
    info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=60).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.M)
    return int(match.group(1)) if match else 0, run.stdout[-500:]


def item(identifier, passed, evidence):
    description, weight, kind = SPECS[identifier]
    return {"criterion_id": identifier, "score": 1.0 if passed else 0.0, "evidence": evidence, "method": "deterministic", "raw": {"description": description, "weight": weight, "type": kind, "passed": bool(passed)}}


def assemble(path: Path, criteria):
    hurdles = [x["criterion_id"] for x in criteria if x["raw"]["type"] == "hurdle" and not x["raw"]["passed"]]
    total = sum(x["raw"]["weight"] for x in criteria if x["raw"]["weight"] > 0)
    earned = sum(x["raw"]["weight"] for x in criteria if x["raw"]["weight"] > 0 and x["raw"]["passed"])
    penalty = sum(abs(x["raw"]["weight"]) for x in criteria if x["raw"]["weight"] < 0 and not x["raw"]["passed"])
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [sha256(path)] if path.is_file() else [], "criterion_results": criteria, "reward": reward, "passed": reward >= 0.8 and not hurdles, "hurdle_failures": hurdles}


def evaluate(path: Path, artifact_dir: Path):
    source_ok, source_evidence = source_check()
    try:
        prs = Presentation(path)
    except Exception as error:
        return assemble(path, [item(identifier, False, f"打开失败：{error}") for identifier in SPECS])
    package_ok, package_evidence = package_checks(path)
    ratio = prs.slide_width / prs.slide_height if prs.slide_height else 0
    r1 = source_ok and package_ok and len(prs.slides) == 6 and abs(ratio - 16 / 9) < 0.02
    criteria = [item("R001", r1, f"页={len(prs.slides)} 比例={ratio:.4f} source={source_evidence} package={json.dumps(package_evidence, ensure_ascii=False)}")]

    charts = chart_mappings(prs)
    revenue = [(slide, mapping) for slide, mapping in charts if mapping == {"Q1 FY23": 8288.0, "Q4 FY23": 6051.0, "Q1 FY24": 7192.0}]
    r2 = len(revenue) == 1
    criteria.append(item("R002", r2, f"Revenue matches={revenue}; all_charts={charts}"))

    tables = table_candidates(prs)
    gaap = find_table(tables, ["Metric", "Q1 FY24", "Q4 FY23", "Q1 FY23", "Q/Q", "Y/Y"], GAAP_ROWS)
    r3 = len(gaap) == 1
    criteria.append(item("R003", r3, f"GAAP table slides={gaap}"))

    segment_chart = [(slide, mapping) for slide, mapping in charts if mapping == {"Data Center": 4.28, "Gaming": 2.24, "Professional Visualization": 0.295, "Automotive": 0.296}]
    segment_slides = [(number, slide_text(slide)) for number, slide in enumerate(prs.slides, 1) if all(token in slide_text(slide) for token in SEGMENT_TOKENS)]
    r4 = len(segment_chart) == 1 and len(segment_slides) == 1 and segment_slides[0][1].count("Record quarter") >= 2
    criteria.append(item("R004", r4, f"segment_chart={segment_chart}; token_slides={[x[0] for x in segment_slides]}"))

    non_gaap = find_table(tables, ["Q1 FY2024", "GAAP", "Non-GAAP"], NON_GAAP_ROWS)
    cash_slides = [number for number, slide in enumerate(prs.slides, 1) if all(token.upper() in slide_text(slide).upper() for token in CASH_TOKENS)]
    disclaimer = [number for number, slide in enumerate(prs.slides, 1) if "Non-GAAP measures" in slide_text(slide) and "do not replace GAAP" in slide_text(slide)]
    r5 = len(non_gaap) == 1 and len(cash_slides) == 1 and len(disclaimer) == 1
    criteria.append(item("R005", r5, f"non_gaap={non_gaap}; cash={cash_slides}; disclaimer={disclaimer}"))

    outlook_slides = [(number, slide_text(slide)) for number, slide in enumerate(prs.slides, 1) if all(token in slide_text(slide) for token in OUTLOOK_TOKENS)]
    r6 = len(outlook_slides) == 1
    criteria.append(item("R006", r6, f"outlook slides={[x[0] for x in outlook_slides]}"))

    text = deck_text(prs)
    citations = text.count("Source: OfficeBench financial_report.pdf")
    prohibited_advice = [token for token in ["BUY NVIDIA", "SELL NVIDIA", "PRICE TARGET", "OUTPERFORM RATING"] if token in text.upper()]
    boundary_ok = citations == 6 and "Derived reading" in text and "NO INVESTMENT RECOMMENDATION" in text and "FORWARD-LOOKING" in text and "not a guarantee of future performance" in text and not prohibited_advice
    r7 = boundary_ok
    criteria.append(item("R007", r7, f"citations={citations}; derived={'Derived reading' in text}; no_reco={'NO INVESTMENT RECOMMENDATION' in text}; forward={'FORWARD-LOOKING' in text}; prohibited_advice={prohibited_advice}"))

    native_tables = sum(shape.has_table for slide in prs.slides for shape in slide.shapes)
    native_charts = sum(shape.shape_type == MSO_SHAPE_TYPE.CHART for slide in prs.slides for shape in slide.shapes)
    full_page_pictures = [f"s{number}:{shape.shape_id}" for number, slide in enumerate(prs.slides, 1) for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.width * shape.height >= prs.slide_width * prs.slide_height * 0.8]
    with zipfile.ZipFile(path) as archive: media = [name for name in archive.namelist() if name.startswith("ppt/media/") and not name.endswith("/")]
    markers = [marker for marker in SUSPICIOUS if marker in text.lower()]
    r8 = native_tables >= 2 and native_charts >= 2 and not full_page_pictures and not media and not markers and not package_evidence.get("macros") and not package_evidence.get("external_targets")
    criteria.append(item("R008", r8, f"tables={native_tables} charts={native_charts} full_page={full_page_pictures} media={media} markers={markers} external={package_evidence.get('external_targets')}"))

    pages = 0; log = ""
    try:
        with tempfile.TemporaryDirectory(prefix="officebench32-render-") as directory: pages, log = render(path, Path(directory))
    except Exception as error: log = str(error)
    delivered = sorted(str(item_path.relative_to(artifact_dir)) for item_path in artifact_dir.rglob("*") if item_path.is_file())
    r9 = pages == 6 and delivered == [OUTPUT_NAME] and package_ok
    criteria.append(item("R009", r9, f"pages={pages} files={delivered} log={log}"))

    severe = not all([r2, r3, r4, r5, r6, r7, r8])
    criteria.append(item("R010", not severe, f"severe={severe}"))
    return assemble(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser(); parser.add_argument("--artifact-dir", type=Path, required=True); parser.add_argument("--result", type=Path, required=True); parser.add_argument("--reward-file", type=Path, required=True); args = parser.parse_args()
    path = args.artifact_dir / OUTPUT_NAME
    output = evaluate(path, args.artifact_dir) if path.is_file() else assemble(path, [item(identifier, False, f"缺失 {OUTPUT_NAME}") for identifier in SPECS])
    args.result.parent.mkdir(parents=True, exist_ok=True); args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(output, ensure_ascii=False, indent=2) + "\n"); args.reward_file.write_text(f"{output['reward']:.6f}\n"); print(json.dumps(output, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
