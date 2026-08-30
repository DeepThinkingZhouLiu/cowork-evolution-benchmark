#!/usr/bin/env python3
"""逐张审计 OfficeBench 3-5 两份发票并制作财务复核简报。"""
from __future__ import annotations
import argparse, hashlib
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HASHES={"Invoice.pdf":"29b4643565c16cbfce84114d92cd3edb11d94ed017f727873ddf4fcf5d11d194","Invoice_2.pdf":"5758d6a85a47563ef5fa718a8a8603106b0c83c4a129f4f91dc57210809df8ec"}
NAVY=RGBColor(0x13,0x20,0x3A); INDIGO=RGBColor(0x4F,0x46,0xE5); CYAN=RGBColor(0x08,0x91,0xB2); ORANGE=RGBColor(0xEA,0x58,0x0C); RED=RGBColor(0xC2,0x31,0x2D); GREEN=RGBColor(0x16,0x8A,0x62); WHITE=RGBColor(0xFF,0xFF,0xFF); INK=RGBColor(0x1F,0x29,0x37); MUTED=RGBColor(0x66,0x70,0x85); LINE=RGBColor(0xE3,0xE7,0xEF); SURFACE=RGBColor(0xF6,0xF7,0xFB)

def sha256(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def add_text(slide,text,left,top,width,height,*,size=15,color=INK,bold=False,align=PP_ALIGN.LEFT,valign=MSO_ANCHOR.TOP,margin=.04):
    shape=slide.shapes.add_textbox(Inches(left),Inches(top),Inches(width),Inches(height)); frame=shape.text_frame; frame.clear(); frame.word_wrap=True; frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=Inches(margin); frame.vertical_anchor=valign; paragraph=frame.paragraphs[0]; paragraph.alignment=align; run=paragraph.add_run(); run.text=str(text); run.font.name="Aptos"; run.font.size=Pt(size); run.font.color.rgb=color; run.font.bold=bold; return shape
def add_rect(slide,left,top,width,height,fill,*,line=None,rounded=True):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(left),Inches(top),Inches(width),Inches(height)); shape.fill.solid(); shape.fill.fore_color.rgb=fill; shape.line.color.rgb=line or fill; return shape
def background(slide,color): slide.background.fill.solid(); slide.background.fill.fore_color.rgb=color
def header(slide,title,kicker,page):
    add_text(slide,kicker.upper(),.65,.28,7.5,.24,size=9,color=INDIGO,bold=True); add_text(slide,title,.65,.57,11.7,.55,size=27,color=NAVY,bold=True); add_text(slide,f"{page:02d} / 05",11.65,.33,.95,.2,size=9,color=MUTED,align=PP_ALIGN.RIGHT); add_rect(slide,.65,1.18,12.03,.02,LINE,rounded=False)
def footer(slide,page,sources,dark=False):
    color=RGBColor(0xAE,0xB8,0xC8) if dark else MUTED; add_text(slide,f"Source: {sources}  ·  OfficeBench 3-5",.65,7.13,10.8,.18,size=7.5,color=color); add_text(slide,str(page),12.15,7.13,.45,.18,size=8,color=color,align=PP_ALIGN.RIGHT)
def metric(slide,value,label,detail,left,top,width,accent=INDIGO,dark=False):
    fill=RGBColor(0x1D,0x2B,0x49) if dark else WHITE; line=RGBColor(0x32,0x43,0x65) if dark else LINE; primary=WHITE if dark else NAVY; secondary=RGBColor(0xB6,0xC2,0xD4) if dark else MUTED
    add_rect(slide,left,top,width,1.12,fill,line=line); add_rect(slide,left,top,.07,1.12,accent,rounded=False); add_text(slide,value,left+.22,top+.13,width-.3,.38,size=22,color=primary,bold=True); add_text(slide,label.upper(),left+.22,top+.58,width-.3,.2,size=8.5,color=secondary,bold=True); add_text(slide,detail,left+.22,top+.83,width-.3,.2,size=8.5,color=accent,bold=True)
def style_table(table,widths,font_size=9):
    for col,width in zip(table.columns,widths,strict=True): col.width=Inches(width)
    for cell in table.rows[0].cells:
        cell.fill.solid(); cell.fill.fore_color.rgb=NAVY
        for run in cell.text_frame.paragraphs[0].runs: run.font.name="Aptos"; run.font.size=Pt(font_size); run.font.bold=True; run.font.color.rgb=WHITE
    for r in range(1,len(table.rows)):
        for cell in table.rows[r].cells:
            cell.fill.solid(); cell.fill.fore_color.rgb=WHITE if r%2 else SURFACE; cell.margin_left=cell.margin_right=Inches(.06)
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs: run.font.name="Aptos"; run.font.size=Pt(font_size); run.font.color.rgb=INK

def build(input_dir:Path,output:Path):
    for name,expected in HASHES.items():
        path=input_dir/name
        if not path.is_file() or sha256(path)!=expected: raise ValueError(f"冻结输入缺失或 SHA256 错：{name}")
    total=104.98+1389.99; subtotal=99.99+1300.00; tax=4.99+65.00; shipping=24.99
    if round(total,2)!=1494.97 or round(subtotal+tax+shipping,2)!=1494.97: raise ValueError("人工审计汇总不一致")
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]

    slide=prs.slides.add_slide(blank); background(slide,NAVY)
    add_text(slide,"AP REVIEW  /  JANUARY 2023",.72,.62,6,.25,size=10,color=RGBColor(0x8B,0xE9,0xF4),bold=True)
    add_text(slide,"Two invoices.\nOne identifier collision.",.72,1.16,7.4,1.32,size=35,color=WHITE,bold=True)
    add_text(slide,"Payment-trend and data-quality briefing",.75,2.82,6.6,.34,size=14,color=RGBColor(0xC7,0xD0,0xDE))
    add_rect(slide,9.17,.95,3.4,2.52,RGBColor(0x1D,0x2B,0x49),line=RGBColor(0x32,0x43,0x65)); add_text(slide,"$1,494.97",9.39,1.38,2.96,.55,size=29,color=WHITE,bold=True,align=PP_ALIGN.CENTER); add_text(slide,"TOTAL DUE",9.39,2.05,2.96,.22,size=9,color=RGBColor(0xB6,0xC2,0xD4),bold=True,align=PP_ALIGN.CENTER); add_text(slide,"2 source invoices",9.39,2.54,2.96,.26,size=12,color=RGBColor(0x8B,0xE9,0xF4),bold=True,align=PP_ALIGN.CENTER)
    metric(slide,"$1,389.99","Largest invoice","93.0% of total due",.72,4.75,3.75,ORANGE,dark=True); metric(slide,"29 days","Invoice-date span","1 Jan -> 30 Jan",4.79,4.75,3.75,CYAN,dark=True); metric(slide,"#100 ×2","Invoice identifier","different source documents",8.86,4.75,3.75,RED,dark=True)
    add_text(slide,"NO PAYMENT-STATUS EVIDENCE IN SOURCE PDFs",.75,6.39,6.5,.22,size=9,color=RGBColor(0xB6,0xC2,0xD4),bold=True); footer(slide,1,"Invoice.pdf p.1; Invoice_2.pdf p.1",dark=True)

    slide=prs.slides.add_slide(blank); background(slide,WHITE); header(slide,"Invoice register", "Extracted source facts",2)
    rows=[
        ["Source","Invoice","Date","Customer","Description","Qty","Unit","Subtotal","Tax","S&H","Total due"],
        ["Invoice.pdf","#100","1/30/23","The Social Strategists","Cloud service","1 TB","$99.99","$99.99","$4.99","$0.00","$104.98"],
        ["Invoice_2.pdf","#100","1/1/23","Perfect Places Interior Design","Decorative clay pottery (LG)","100","$13.00","$1,300.00","$65.00","$24.99","$1,389.99"],
    ]
    table=slide.shapes.add_table(3,11,Inches(.52),Inches(1.55),Inches(12.28),Inches(2.05)).table
    for r,values in enumerate(rows):
        for c,value in enumerate(values): table.cell(r,c).text=value
    style_table(table,[1.1,.65,.73,1.62,1.72,.55,.75,.92,.7,.7,.94],font_size=7.5)
    for r in range(3): table.rows[r].height=Inches(.55 if r==0 else .75)
    add_rect(slide,.72,4.05,5.78,1.65,SURFACE,line=LINE); add_text(slide,"CLOUD SERVICE INVOICE",1.02,4.35,2.8,.24,size=10,color=INDIGO,bold=True); add_text(slide,"Purchased by / ship to: Liane Cormier\nTerms: Due upon receipt\nVendor name: not printed in the PDF",1.02,4.76,4.95,.72,size=12,color=INK)
    add_rect(slide,6.82,4.05,5.78,1.65,SURFACE,line=LINE); add_text(slide,"POTTERY INVOICE",7.12,4.35,2.8,.24,size=10,color=ORANGE,bold=True); add_text(slide,"Bill / ship to: Mollie Grau\nTerms: Due on receipt  ·  Express air\nPayee: Pottery & Co.  ·  Fragile goods",7.12,4.76,4.95,.72,size=12,color=INK)
    add_text(slide,"Both PDFs display invoice #100; the source filename and customer must remain part of any unique review key.",.78,6.22,11.8,.36,size=11,color=RED,bold=True); footer(slide,2,"Invoice.pdf p.1; Invoice_2.pdf p.1")

    slide=prs.slides.add_slide(blank); background(slide,SURFACE); header(slide,"Spend is concentrated in pottery", "Composition",3)
    data=ChartData(); data.categories=["Cloud service","Decorative clay pottery"]; data.add_series("Subtotal",[99.99,1300.00]); data.add_series("Sales tax",[4.99,65.00]); data.add_series("Shipping & handling",[0.00,24.99])
    chart=slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED,Inches(.72),Inches(1.58),Inches(7.05),Inches(4.9),data).chart; chart.has_title=False; chart.has_legend=True; chart.legend.include_in_layout=False; chart.value_axis.minimum_scale=0; chart.value_axis.maximum_scale=1500
    colors=[INDIGO,ORANGE,CYAN]
    for series,color in zip(chart.series,colors,strict=True): series.format.fill.solid(); series.format.fill.fore_color.rgb=color
    plot=chart.plots[0]; plot.has_data_labels=True; plot.data_labels.show_value=True; plot.data_labels.position=XL_DATA_LABEL_POSITION.CENTER
    metric(slide,"$1,399.99","Aggregate subtotal","before tax and shipping",8.13,1.58,4.28,INDIGO); metric(slide,"$69.99","Aggregate sales tax","≈5.00% of subtotal",8.13,2.93,4.28,ORANGE); metric(slide,"$24.99","Aggregate shipping","pottery invoice only",8.13,4.28,4.28,CYAN)
    add_text(slide,"Pottery invoice share of total due: 93.0%  ·  Cloud invoice share: 7.0%",8.18,5.81,4.14,.52,size=11,color=NAVY,bold=True); footer(slide,3,"Invoice.pdf p.1; Invoice_2.pdf p.1")

    slide=prs.slides.add_slide(blank); background(slide,WHITE); header(slide,"Payment trend: terms, not status", "Control interpretation",4)
    metric(slide,"2 / 2","Payment terms","due upon receipt",.72,1.55,3.75,INDIGO); metric(slide,"0 / 2","Payment records","none supplied",4.79,1.55,3.75,RED); metric(slide,"29 days","Document cadence","invoice-date span",8.86,1.55,3.75,CYAN)
    rows=[["Source","Date","Terms","Sales-tax rate","Shipping","Payment status"],["Invoice.pdf","1/30/23","Due upon receipt","4.99 / 99.99 ≈ 4.99%","$0.00","Not provided"],["Invoice_2.pdf","1/1/23","Due on receipt","65.00 / 1,300 = 5.00%","$24.99","Not provided"]]
    table=slide.shapes.add_table(3,6,Inches(.72),Inches(3.1),Inches(11.9),Inches(2.15)).table
    for r,values in enumerate(rows):
        for c,value in enumerate(values): table.cell(r,c).text=value
    style_table(table,[1.45,1.05,1.75,2.65,1.2,2.25],font_size=9.5)
    for r in range(3): table.rows[r].height=Inches(.56 if r==0 else .78)
    add_rect(slide,.72,5.73,11.9,.86,RGBColor(0xFF,0xF5,0xF2),line=RGBColor(0xFD,0xC8,0xB8)); add_text(slide,"Do not label either invoice paid, unpaid, late or overdue: the PDFs contain terms and dates, but no payment event, due-date calculation or review date.",1.0,5.98,11.35,.38,size=11,color=RED,bold=True); footer(slide,4,"Invoice.pdf p.1; Invoice_2.pdf p.1")

    slide=prs.slides.add_slide(blank); background(slide,NAVY)
    add_text(slide,"FINANCIAL REVIEW  /  NEXT CONTROLS",.72,.58,7.4,.24,size=10,color=RGBColor(0x8B,0xE9,0xF4),bold=True); add_text(slide,"Resolve identity before payment",.72,1.02,7.6,.58,size=31,color=WHITE,bold=True)
    actions=[("1","Disambiguate invoice #100","Use source filename + customer + date; do not deduplicate on invoice number alone.",RED),("2","Verify the unnamed cloud vendor","Invoice.pdf prints an address and phone but no vendor/company name.",ORANGE),("3","Validate pottery logistics","Match 100 units, fragile-goods note, Express air and $24.99 shipping to receipt evidence.",CYAN),("4","Obtain payment evidence","Add approval, payment date, payment method and clearing status before trend claims.",INDIGO)]
    for index,(num,title,detail,accent) in enumerate(actions):
        top=1.92+index*1.07; add_rect(slide,.72,top,8.0,.86,RGBColor(0x1D,0x2B,0x49),line=RGBColor(0x32,0x43,0x65)); add_text(slide,num,.96,top+.18,.35,.34,size=20,color=accent,bold=True); add_text(slide,title,1.46,top+.13,2.75,.28,size=14,color=WHITE,bold=True); add_text(slide,detail,4.22,top+.13,4.12,.5,size=10,color=RGBColor(0xC7,0xD0,0xDE))
    add_rect(slide,9.1,1.92,3.45,4.08,WHITE,line=WHITE); add_text(slide,"REVIEW BOUNDARY",9.43,2.25,2.75,.24,size=9,color=INDIGO,bold=True); add_text(slide,"Known",9.43,2.78,2.4,.27,size=14,color=NAVY,bold=True); add_text(slide,"Dates, terms, line items, tax, shipping and total due",9.43,3.13,2.45,.68,size=11,color=INK); add_text(slide,"Unknown",9.43,4.07,2.4,.27,size=14,color=RED,bold=True); add_text(slide,"Paid / unpaid, overdue, payment date, approval and settlement method",9.43,4.42,2.45,.85,size=11,color=INK); add_text(slide,"No external assumptions",9.43,5.53,2.45,.24,size=10,color=GREEN,bold=True)
    footer(slide,5,"Invoice.pdf p.1; Invoice_2.pdf p.1",dark=True)
    output.parent.mkdir(parents=True,exist_ok=True); prs.save(output)
    if len(Presentation(output).slides)!=5: raise ValueError("Golden 不是 5 页")

def main():
    p=argparse.ArgumentParser(); p.add_argument("--input-dir",type=Path,required=True); p.add_argument("--output",type=Path,required=True); a=p.parse_args(); build(a.input_dir,a.output); print(f"{a.output}\nsha256={sha256(a.output)}")
if __name__=="__main__": main()
