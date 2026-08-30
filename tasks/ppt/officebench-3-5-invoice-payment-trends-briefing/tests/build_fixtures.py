#!/usr/bin/env python3
from __future__ import annotations
import argparse,shutil,zipfile
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches
OUTPUT_NAME="January_2023_Invoice_Payment_Review.pptx"; C="http://schemas.openxmlformats.org/drawingml/2006/chart"
def out(root,name): p=root/name/OUTPUT_NAME; p.parent.mkdir(parents=True,exist_ok=True); return p
def mutate(g,o,fn): p=Presentation(g); fn(p); p.save(o)
def replace_all(prs,old,new):
    n=0
    for sl in prs.slides:
        for s in sl.shapes:
            if getattr(s,"has_text_frame",False) and old in s.text: s.text=s.text.replace(old,new); n+=1
    if not n: raise ValueError(old)
def table(prs,header):
    for sl in prs.slides:
        for s in sl.shapes:
            if s.has_table and s.table.cell(0,0).text==header: return s.table
    raise ValueError(header)
def main():
    p=argparse.ArgumentParser(); p.add_argument("--golden",type=Path,required=True); p.add_argument("--outdir",type=Path,required=True); a=p.parse_args(); (a.outdir/"empty").mkdir(parents=True); out(a.outdir,"corrupt").write_bytes(b"bad\n")
    with zipfile.ZipFile(a.golden) as r,zipfile.ZipFile(out(a.outdir,"valid-repackage"),"w",compression=zipfile.ZIP_STORED) as w:
        for i in r.infolist(): w.writestr(i.filename,r.read(i.filename))
    def alt(prs): prs.slides[0].background.fill.solid(); prs.slides[0].background.fill.fore_color.rgb=RGBColor(0x18,0x28,0x44)
    mutate(a.golden,out(a.outdir,"valid-alt-cover"),alt)
    def missing(prs):
        t=table(prs,"Source")
        for c in t.rows[2].cells: c.text=""
    mutate(a.golden,out(a.outdir,"missing-invoice-row"),missing)
    mutate(a.golden,out(a.outdir,"wrong-invoice-total"),lambda prs: setattr(table(prs,"Source").cell(2,10),"text","$1,399.99"))
    def wrong_chart(prs):
        c=next(s.chart for sl in prs.slides for s in sl.shapes if s.shape_type==MSO_SHAPE_TYPE.CHART); cache=c._chartSpace.find(f".//{{{C}}}numCache"); cache.findall(f"{{{C}}}pt")[0].find(f"{{{C}}}v").text="199.99"
    mutate(a.golden,out(a.outdir,"wrong-composition-chart"),wrong_chart)
    mutate(a.golden,out(a.outdir,"wrong-aggregate"),lambda prs: replace_all(prs,"$1,494.97","$1,500.00"))
    mutate(a.golden,out(a.outdir,"wrong-date-span"),lambda prs: replace_all(prs,"29 days","28 days"))
    def paid(prs):
        for sl in prs.slides:
            for s in sl.shapes:
                if s.has_table and len(s.table.columns)==6 and s.table.cell(0,0).text=="Source":
                    s.table.cell(1,5).text="Paid"; return
        raise ValueError("payment table")
    mutate(a.golden,out(a.outdir,"invented-payment-status"),paid)
    mutate(a.golden,out(a.outdir,"invented-cloud-vendor"),lambda prs: replace_all(prs,"Vendor name: not printed in the PDF","Vendor: CloudCo"))
    mutate(a.golden,out(a.outdir,"wrong-pottery-logistics"),lambda prs: replace_all(prs,"Express air","Ground"))
    mutate(a.golden,out(a.outdir,"missing-citations"),lambda prs: replace_all(prs,"Source: Invoice.pdf p.1; Invoice_2.pdf p.1","Sources omitted"))
    image=a.outdir/"flat.png"; Image.new("RGB",(1600,900),"white").save(image)
    def flat(prs):
        sl=prs.slides[1]
        for s in list(sl.shapes): s._element.getparent().remove(s._element)
        sl.shapes.add_picture(str(image),0,0,width=prs.slide_width,height=prs.slide_height)
    mutate(a.golden,out(a.outdir,"flattened-register"),flat)
    def marker(prs): b=prs.slides[-1].shapes.add_textbox(Inches(.1),Inches(.1),Inches(2),Inches(.2)); b.text="solution/Golden"
    mutate(a.golden,out(a.outdir,"hardcoded-marker"),marker)
    extra=a.outdir/"extra-output"; extra.mkdir(); shutil.copy2(a.golden,extra/OUTPUT_NAME); (extra/"debug.txt").write_text("debug\n")
if __name__=="__main__": main()
