#!/usr/bin/env python3
"""OfficeBench 3-5 发票复核简报的任务专属 Judge。"""
from __future__ import annotations
import argparse,hashlib,json,posixpath,re,subprocess,tempfile,zipfile
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

OUTPUT_NAME="January_2023_Invoice_Payment_Review.pptx"; TASK_ID="cowork-evo/officebench-3-5-invoice-payment-trends-briefing"; RUBRIC_VERSION="officebench-3-5-ppt-adaptation-v0.1"; JUDGE_VERSION="officebench-3-5-deterministic-v0.1"
HASHES={"Invoice.pdf":"29b4643565c16cbfce84114d92cd3edb11d94ed017f727873ddf4fcf5d11d194","Invoice_2.pdf":"5758d6a85a47563ef5fa718a8a8603106b0c83c4a129f4f91dc57210809df8ec"}
REGISTER_HEADERS=["Source","Invoice","Date","Customer","Description","Qty","Unit","Subtotal","Tax","S&H","Total due"]
REGISTER_ROWS={"Invoice.pdf":("#100","1/30/23","The Social Strategists","Cloud service","1 TB","$99.99","$99.99","$4.99","$0.00","$104.98"),"Invoice_2.pdf":("#100","1/1/23","Perfect Places Interior Design","Decorative clay pottery (LG)","100","$13.00","$1,300.00","$65.00","$24.99","$1,389.99")}
PAY_HEADERS=["Source","Date","Terms","Sales-tax rate","Shipping","Payment status"]
PAY_ROWS={"Invoice.pdf":("1/30/23","Due upon receipt","4.99 / 99.99 ≈ 4.99%","$0.00","Not provided"),"Invoice_2.pdf":("1/1/23","Due on receipt","65.00 / 1,300 = 5.00%","$24.99","Not provided")}
SUMMARY_TOKENS=["$1,399.99","$69.99","$24.99","$1,494.97","93.0%","7.0%","#100 ×2","29 days"]
CONTROL_TOKENS=["Vendor name: not printed","100 units","fragile-goods note","Express air","$24.99 shipping","Pottery & Co.","Disambiguate invoice #100","Verify the unnamed cloud vendor","Validate pottery logistics","Obtain payment evidence","No external assumptions"]
BOUNDARY_TOKENS=["Payment trend: terms, not status","0 / 2","PAYMENT RECORDS","Do not label either invoice paid, unpaid, late or overdue","no payment event"]
SUSPICIOUS=["solution/golden","judge-result","artifact_dir","officebench-3-5","/root/"]
SPECS={"R001":("PPTX 完整、16:9 且 5 页",6,"hurdle"),"R002":("两行发票登记准确",24,"hurdle"),"R003":("原生金额构成图准确",16,"hurdle"),"R004":("汇总、构成、编号和跨度准确",14,"hurdle"),"R005":("条款与付款状态边界准确",16,"hurdle"),"R006":("来源问题与复核动作完整",10,"hurdle"),"R007":("逐页来源且无外部事实",4,"hurdle"),"R008":("原生可编辑且无标记/宏/外链",8,"hurdle"),"R009":("5 页可渲染且输出干净",2,"positive"),"R010":("无严重事实或边界错误",-20,"penalty")}
def sha256(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def norm(v): return re.sub(r"\s+"," ",str(v).replace("\u00a0"," ").replace("\v"," ")).strip()
def locate(name):
    for p in (Path("/app/input_files")/name,Path(__file__).resolve().parents[1]/"data/input_files"/name):
        if p.is_file(): return p
    return None
def source_check():
    missing=[]
    expected_tokens={"Invoice.pdf":["INVOICE # 100","DATE: 1/30/23","Cloud service","99.99","4.99","104.98","Due upon receipt","The Social Strategists"],"Invoice_2.pdf":["INVOICE #100","DATE: 1/1/23","Decorative clay pottery (LG)","1300.00","65.00","24.99","1389.99","Pottery & Co.","Due on receipt"]}
    for name,expected in HASHES.items():
        p=locate(name)
        if p is None or sha256(p)!=expected: missing.append(f"{name}:hash"); continue
        text=subprocess.run(["pdftotext","-layout",str(p),"-"],stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,timeout=60).stdout
        missing.extend(f"{name}:{token}" for token in expected_tokens[name] if token not in text)
    return not missing,missing
def package_checks(path):
    e={"broken":None,"missing":[],"macros":[],"external":[]}
    try:
        with zipfile.ZipFile(path) as z:
            e["broken"]=z.testzip(); names=set(z.namelist()); e["macros"]=[n for n in names if n.lower().endswith("vbaproject.bin")]
            for rn in sorted(n for n in names if n.endswith(".rels")):
                root=etree.fromstring(z.read(rn)); owner="" if rn=="_rels/.rels" else posixpath.join(rn.split("/_rels/",1)[0],rn.split("/_rels/",1)[1][:-5]); base=posixpath.dirname(owner)
                for rel in root:
                    target=rel.attrib.get("Target","")
                    if rel.attrib.get("TargetMode")=="External": e["external"].append(target); continue
                    resolved=target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base,target))
                    if resolved not in names: e["missing"].append(f"{rn}->{resolved}")
        return e["broken"] is None and not e["missing"] and not e["macros"],e
    except Exception as x: e["exception"]=str(x); return False,e
def slide_text(slide): return "\n".join(s.text for s in slide.shapes if getattr(s,"has_text_frame",False) and not s.has_table)
def deck_text(prs): return "\n".join(s.text for sl in prs.slides for s in sl.shapes if getattr(s,"has_text_frame",False))
def tables(prs):
    out=[]
    for sn,sl in enumerate(prs.slides,1):
        for s in sl.shapes:
            if s.has_table: out.append((sn,[[norm(c.text) for c in row.cells] for row in s.table.rows]))
    return out
def exact_table(candidates,headers,expected):
    return [sn for sn,rows in candidates if rows and rows[0]==headers and {row[0]:tuple(row[1:]) for row in rows[1:] if row and row[0]}==expected]
def charts(prs):
    result=[]
    for sn,sl in enumerate(prs.slides,1):
        for s in sl.shapes:
            if s.shape_type!=MSO_SHAPE_TYPE.CHART: continue
            c=s.chart; cats=[norm(x.label) for x in c.plots[0].categories]; series=[]
            for ser in c.series: series.append((norm(ser.name),dict(zip(cats,[float(v) for v in ser.values],strict=True))))
            result.append((sn,series))
    return result
def render(path,out):
    out.mkdir(parents=True,exist_ok=True); profile=out/"profile"; profile.mkdir(exist_ok=True); run=subprocess.run(["soffice",f"-env:UserInstallation=file://{profile}","--headless","--convert-to","pdf","--outdir",str(out),str(path)],stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,timeout=300); pdf=out/f"{path.stem}.pdf"
    if run.returncode or not pdf.is_file(): return 0,run.stdout[-500:]
    info=subprocess.run(["pdfinfo",str(pdf)],stdout=subprocess.PIPE,text=True).stdout; m=re.search(r"^Pages:\s+(\d+)",info,re.M); return int(m.group(1)) if m else 0,run.stdout[-500:]
def item(i,p,e): d,w,k=SPECS[i]; return {"criterion_id":i,"score":1.0 if p else 0.0,"evidence":e,"method":"deterministic","raw":{"description":d,"weight":w,"type":k,"passed":bool(p)}}
def assemble(path,criteria):
    hurdles=[x["criterion_id"] for x in criteria if x["raw"]["type"]=="hurdle" and not x["raw"]["passed"]]; total=sum(x["raw"]["weight"] for x in criteria if x["raw"]["weight"]>0); earned=sum(x["raw"]["weight"] for x in criteria if x["raw"]["weight"]>0 and x["raw"]["passed"]); penalty=sum(abs(x["raw"]["weight"]) for x in criteria if x["raw"]["weight"]<0 and not x["raw"]["passed"]); reward=0.0 if hurdles else max(0,min(1,(earned-penalty)/total)); return {"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":[sha256(path)] if path.is_file() else [],"criterion_results":criteria,"reward":reward,"passed":reward>=.8 and not hurdles,"hurdle_failures":hurdles}
def evaluate(path,artifact):
    source_ok,se=source_check()
    try: prs=Presentation(path)
    except Exception as x: return assemble(path,[item(i,False,f"打开失败：{x}") for i in SPECS])
    pkg,pe=package_checks(path); ratio=prs.slide_width/prs.slide_height if prs.slide_height else 0; r1=source_ok and pkg and len(prs.slides)==5 and abs(ratio-16/9)<.02; criteria=[item("R001",r1,f"pages={len(prs.slides)} ratio={ratio:.4f} source={se} package={json.dumps(pe,ensure_ascii=False)}")]
    tc=tables(prs); register=exact_table(tc,REGISTER_HEADERS,REGISTER_ROWS); r2=len(register)==1; criteria.append(item("R002",r2,f"register={register}"))
    observed=charts(prs); expected=[("Subtotal",{"Cloud service":99.99,"Decorative clay pottery":1300.0}),("Sales tax",{"Cloud service":4.99,"Decorative clay pottery":65.0}),("Shipping & handling",{"Cloud service":0.0,"Decorative clay pottery":24.99})]; chart_matches=[sn for sn,series in observed if series==expected]; r3=len(chart_matches)==1; criteria.append(item("R003",r3,f"chart_matches={chart_matches}; charts={observed}"))
    text=deck_text(prs); r4=all(t in text for t in SUMMARY_TOKENS) and text.count("#100")>=3; criteria.append(item("R004",r4,f"missing={[t for t in SUMMARY_TOKENS if t not in text]} #100_count={text.count('#100')}"))
    pay=exact_table(tc,PAY_HEADERS,PAY_ROWS); boundary_slides=[sn for sn,sl in enumerate(prs.slides,1) if all(t.upper() in slide_text(sl).upper() for t in BOUNDARY_TOKENS)]; forbidden=[p for p in ["STATUS: PAID","STATUS: UNPAID","OVERDUE BY","BOTH INVOICES ARE PAID","BOTH INVOICES ARE UNPAID"] if p in text.upper()]; r5=len(pay)==1 and len(boundary_slides)==1 and not forbidden; criteria.append(item("R005",r5,f"payment_table={pay} boundary={boundary_slides} forbidden={forbidden}"))
    r6=all(t.upper() in text.upper() for t in CONTROL_TOKENS); criteria.append(item("R006",r6,f"missing={[t for t in CONTROL_TOKENS if t.upper() not in text.upper()]}"))
    citations=text.count("Source: Invoice.pdf p.1; Invoice_2.pdf p.1"); r7=citations==5 and "No external assumptions" in text; criteria.append(item("R007",r7,f"citations={citations}"))
    native_tables=sum(s.has_table for sl in prs.slides for s in sl.shapes); native_charts=sum(s.shape_type==MSO_SHAPE_TYPE.CHART for sl in prs.slides for s in sl.shapes); full=[f"s{sn}:{s.shape_id}" for sn,sl in enumerate(prs.slides,1) for s in sl.shapes if s.shape_type==MSO_SHAPE_TYPE.PICTURE and s.width*s.height>=prs.slide_width*prs.slide_height*.8]
    with zipfile.ZipFile(path) as z: media=[n for n in z.namelist() if n.startswith("ppt/media/") and not n.endswith("/")]
    markers=[m for m in SUSPICIOUS if m in text.lower()]; r8=native_tables>=2 and native_charts>=1 and not full and not media and not markers and not pe.get("macros") and not pe.get("external"); criteria.append(item("R008",r8,f"tables={native_tables} charts={native_charts} full={full} media={media} markers={markers} external={pe.get('external')}"))
    pages=0; log=""
    try:
        with tempfile.TemporaryDirectory(prefix="officebench35-render-") as d: pages,log=render(path,Path(d))
    except Exception as x: log=str(x)
    delivered=sorted(str(p.relative_to(artifact)) for p in artifact.rglob("*") if p.is_file()); r9=pages==5 and delivered==[OUTPUT_NAME] and pkg; criteria.append(item("R009",r9,f"pages={pages} files={delivered} log={log}")); severe=not all([r2,r3,r4,r5,r6,r7,r8]); criteria.append(item("R010",not severe,f"severe={severe}")); return assemble(path,criteria)
def main():
    p=argparse.ArgumentParser(); p.add_argument("--artifact-dir",type=Path,required=True); p.add_argument("--result",type=Path,required=True); p.add_argument("--reward-file",type=Path,required=True); a=p.parse_args(); path=a.artifact_dir/OUTPUT_NAME; result=evaluate(path,a.artifact_dir) if path.is_file() else assemble(path,[item(i,False,f"缺失 {OUTPUT_NAME}") for i in SPECS]); a.result.parent.mkdir(parents=True,exist_ok=True); a.reward_file.parent.mkdir(parents=True,exist_ok=True); a.result.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n"); a.reward_file.write_text(f"{result['reward']:.6f}\n"); print(json.dumps(result,ensure_ascii=False,indent=2))
if __name__=="__main__": main()
