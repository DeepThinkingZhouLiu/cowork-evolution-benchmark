#!/usr/bin/env python3
"""逐页完成 OmegaUse-OfficeVal officeval_041 的真实 PPT 备注修改任务。"""

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Inches, Pt


OUTPUT_NAME = "高速成型线碎屑分离与状态监测装置_修改完成版.pptx"
SOURCE_NAME = "高速成型线碎屑分离与状态监测装置.pptx"
WORKBOOK_NAME = "对比分析表.xlsx"
IMAGE_NAMES = ["image_001.png", "image_002.png", "image_003.png"]

INPUT_HASHES = {
    "image_001.png": "cfdaf30c5c63be8d72d2e3a67e155bbdd3c46913a5784ca5b3bf61852c150150",
    "image_002.png": "57de004c312ce1a7dd6eaee93bcdf876a8f1d9ee0aecf875dcf16571463546a8",
    "image_003.png": "88573eede478416ac4226dedc802d804fc29a5592e022f214464bf8915638b51",
    "对比分析表.xlsx": "fb76919c552ed81af7abbda44a740c47001e35386e35f9a6d7fa3056e41f5eff",
    "高速成型线碎屑分离与状态监测装置.pptx": "c71a4b61fdc11ae33e77838be01fb7e034791d744678157f19069037e8df7d70",
}

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NSMAP = {"p": P_NS, "a": A_NS, "r": R_NS}

YAHEI = "微软雅黑"
TEAL = RGBColor(0x00, 0xA6, 0xA6)
DEEP_BLUE = RGBColor(0x33, 0x41, 0x55)
SLATE = RGBColor(0x64, 0x74, 0x8B)
NAVY = RGBColor(0x0A, 0x1F, 0x33)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
LIGHT_BLUE = RGBColor(0x93, 0xC5, 0xFD)


def sha256(path: Path) -> str:
    """计算文件 SHA256。"""
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_inputs(input_dir: Path) -> None:
    """保证回放只使用上游发布的五个真实输入文件。"""
    errors = []
    for name, expected in INPUT_HASHES.items():
        path = input_dir / name
        if not path.is_file():
            errors.append(f"缺少输入文件：{name}")
            continue
        actual = sha256(path)
        if actual != expected:
            errors.append(f"输入文件哈希不一致：{name}，实际 {actual}")
    if errors:
        raise ValueError("；".join(errors))


def set_east_asian_font(run, name: str = YAHEI) -> None:
    """同时写入拉丁和东亚字体，避免 Office 回退到主题字体。"""
    run.font.name = name
    r_pr = run._r.get_or_add_rPr()
    ea = r_pr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(r_pr, qn("a:ea"))
    ea.set("typeface", name)


def format_text_frame(
    frame,
    *,
    size: float,
    color: RGBColor,
    font: str = YAHEI,
    bold: bool | None = None,
    align: PP_ALIGN | None = None,
    line_spacing: float | None = None,
) -> None:
    """格式化文本框内所有可见 run。"""
    frame.word_wrap = True
    for paragraph in frame.paragraphs:
        if align is not None:
            paragraph.alignment = align
        if line_spacing is not None:
            paragraph.line_spacing = line_spacing
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        for run in paragraph.runs:
            if not run.text:
                continue
            set_east_asian_font(run, font)
            run.font.size = Pt(size)
            run.font.color.rgb = color
            if bold is not None:
                run.font.bold = bold


def replace_text(
    shape,
    text: str,
    *,
    size: float,
    color: RGBColor,
    font: str = YAHEI,
    bold: bool | None = None,
    align: PP_ALIGN | None = None,
    line_spacing: float | None = None,
) -> None:
    """以一个或多个可编辑段落替换文本，并应用明确格式。"""
    frame = shape.text_frame
    frame.clear()
    lines = text.split("\n")
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = line
    format_text_frame(
        frame,
        size=size,
        color=color,
        font=font,
        bold=bold,
        align=align,
        line_spacing=line_spacing,
    )


def shape_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return shape.text_frame.text.strip()


def find_text_shape(slide, text: str, *, exact: bool = False):
    """定位唯一的原生文本对象。"""
    matches = []
    for shape in slide.shapes:
        value = shape_text(shape)
        if (value == text) if exact else (text in value):
            matches.append(shape)
    if len(matches) != 1:
        raise ValueError(f"PPT 页面内无法唯一定位文本 {text!r}，命中 {len(matches)} 个对象")
    return matches[0]


def remove_shape(shape) -> None:
    """删除一个原生对象。"""
    element = shape._element
    element.getparent().remove(element)


def format_table(
    table_shape,
    *,
    size: float,
    color: RGBColor,
    align: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    """保持表格可编辑，并统一单元格文字格式。"""
    for row in table_shape.table.rows:
        for cell in row.cells:
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Cm(0.12)
            cell.margin_right = Cm(0.12)
            cell.margin_top = Cm(0.05)
            cell.margin_bottom = Cm(0.05)
            format_text_frame(
                cell.text_frame,
                size=size,
                color=color,
                align=align,
                line_spacing=0.95,
            )


def set_cell_border(cell, color: str = "D7E3EA", width: str = "9525") -> None:
    """为新建的 Excel 数据表格写入细边框。"""
    tc_pr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tc_pr.find(qn(tag))
        if old is not None:
            tc_pr.remove(old)
        line = etree.SubElement(tc_pr, qn(tag), w=width)
        solid = etree.SubElement(line, qn("a:solidFill"))
        etree.SubElement(solid, qn("a:srgbClr"), val=color)
        etree.SubElement(line, qn("a:prstDash"), val="solid")


def replace_excel_table(slide, workbook_path: Path) -> None:
    """用真实 Excel 的 Sheet1!A1:D6 替换第 19 页表格。"""
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    worksheet = workbook["Sheet1"]
    values = [
        [worksheet.cell(row=row, column=column).value for column in range(1, 5)]
        for row in range(1, 7)
    ]
    workbook.close()

    old_tables = [shape for shape in slide.shapes if shape.has_table]
    if len(old_tables) != 1:
        raise ValueError(f"第 19 页原始表格数量应为 1，实际为 {len(old_tables)}")
    remove_shape(old_tables[0])

    table_shape = slide.shapes.add_table(
        6,
        4,
        Cm(1.7),
        Cm(3.10),
        Cm(24.07),
        Cm(10.59),
    )
    table = table_shape.table
    column_widths = [Cm(3.30), Cm(6.90), Cm(6.90), Cm(6.97)]
    row_heights = [Cm(1.50)] + [Cm(1.818)] * 5
    for column, width in zip(table.columns, column_widths):
        column.width = width
    for row, height in zip(table.rows, row_heights):
        row.height = height

    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            value = values[row_index][column_index]
            cell.text = "" if value is None else str(value)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Cm(0.10)
            cell.margin_right = Cm(0.10)
            cell.margin_top = Cm(0.04)
            cell.margin_bottom = Cm(0.04)
            if row_index == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0x0F, 0x76, 0x77)
                font_color = WHITE
                bold = True
            elif column_index == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xD9, 0xF0, 0xEF)
                font_color = NAVY
                bold = True
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = (
                    RGBColor(0xF7, 0xFA, 0xFC)
                    if row_index % 2
                    else RGBColor(0xEE, 0xF5, 0xF7)
                )
                font_color = DEEP_BLUE
                bold = False
            format_text_frame(
                cell.text_frame,
                size=9.5,
                color=font_color,
                bold=bold,
                align=PP_ALIGN.CENTER,
                line_spacing=0.9,
            )
            set_cell_border(cell)

    # 指定尺寸必须由列宽、行高之和共同决定。
    table_shape.width = Cm(24.07)
    table_shape.height = Cm(10.59)

    # 保留原页面的现场图片，但缩到数据表右侧，避免遮挡新表格。
    pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    if len(pictures) == 1:
        picture = pictures[0]
        picture.left = Inches(10.50)
        picture.top = Inches(1.30)
        picture.width = Inches(1.78)
        picture.height = Inches(1.50)

    # 表格高度增加后，把原结果框整体下移，保持留白和可读性。
    result_label = find_text_shape(slide, "结果显示：")
    result_text = find_text_shape(slide, "平均分离率约89.8%")
    result_label.top = Inches(5.77)
    result_text.top = Inches(5.69)
    for shape in slide.shapes:
        if shape.has_text_frame or shape.has_table or shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            continue
        if abs(shape.top - Inches(5.25)) < Inches(0.05) and shape.width > Inches(10):
            shape.top = Inches(5.58)
            break


def set_shape_fill_and_line(shape, color: RGBColor) -> None:
    """同时设置圆形填充和轮廓色。"""
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color
    shape.line.width = Pt(1.25)


def find_circle_above(slide, label: str):
    """找到指定标签正上方的正圆。"""
    text_shape = find_text_shape(slide, label, exact=True)
    center_x = text_shape.left + text_shape.width / 2
    candidates = []
    for shape in slide.shapes:
        if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE:
            continue
        geometry = shape._element.spPr.prstGeom
        if geometry is None or geometry.get("prst") != "ellipse":
            continue
        if abs(shape.width - shape.height) > 1000 or shape.top >= text_shape.top:
            continue
        candidates.append((abs((shape.left + shape.width / 2) - center_x), shape))
    if not candidates:
        raise ValueError(f"未找到 {label} 上方圆形")
    return min(candidates, key=lambda item: item[0])[1]


def animation_behavior(parent, timing_id: int, shape_id: str, effect: str | None) -> int:
    """写入显示行为，并按需添加指定进入效果。"""
    child_list = etree.SubElement(parent, f"{{{P_NS}}}childTnLst")
    set_node = etree.SubElement(child_list, f"{{{P_NS}}}set")
    behavior = etree.SubElement(set_node, f"{{{P_NS}}}cBhvr")
    etree.SubElement(
        behavior,
        f"{{{P_NS}}}cTn",
        id=str(timing_id),
        dur="1",
        fill="hold",
    )
    timing_id += 1
    target = etree.SubElement(behavior, f"{{{P_NS}}}tgtEl")
    etree.SubElement(target, f"{{{P_NS}}}spTgt", spid=shape_id)
    attributes = etree.SubElement(behavior, f"{{{P_NS}}}attrNameLst")
    attribute = etree.SubElement(attributes, f"{{{P_NS}}}attrName")
    attribute.text = "style.visibility"
    destination = etree.SubElement(set_node, f"{{{P_NS}}}to")
    etree.SubElement(destination, f"{{{P_NS}}}strVal", val="visible")

    if effect:
        effect_node = etree.SubElement(
            child_list,
            f"{{{P_NS}}}animEffect",
            transition="in",
            filter=effect,
        )
        effect_behavior = etree.SubElement(effect_node, f"{{{P_NS}}}cBhvr")
        etree.SubElement(
            effect_behavior,
            f"{{{P_NS}}}cTn",
            id=str(timing_id),
            dur="650",
        )
        timing_id += 1
        effect_target = etree.SubElement(effect_behavior, f"{{{P_NS}}}tgtEl")
        etree.SubElement(effect_target, f"{{{P_NS}}}spTgt", spid=shape_id)
    return timing_id


def build_single_effect_timing(
    shape_id: int,
    *,
    preset_id: int,
    effect: str | None,
) -> etree._Element:
    """构造一个作用于指定图片的单击进入动画。"""
    timing = etree.Element(f"{{{P_NS}}}timing", nsmap=NSMAP)
    time_nodes = etree.SubElement(timing, f"{{{P_NS}}}tnLst")
    root_parallel = etree.SubElement(time_nodes, f"{{{P_NS}}}par")
    root_ctn = etree.SubElement(
        root_parallel,
        f"{{{P_NS}}}cTn",
        id="1",
        dur="indefinite",
        restart="never",
        nodeType="tmRoot",
    )
    root_children = etree.SubElement(root_ctn, f"{{{P_NS}}}childTnLst")
    sequence = etree.SubElement(root_children, f"{{{P_NS}}}seq", concurrent="1", nextAc="seek")
    main_ctn = etree.SubElement(
        sequence,
        f"{{{P_NS}}}cTn",
        id="2",
        dur="indefinite",
        nodeType="mainSeq",
    )
    main_children = etree.SubElement(main_ctn, f"{{{P_NS}}}childTnLst")
    top_parallel = etree.SubElement(main_children, f"{{{P_NS}}}par")
    top_ctn = etree.SubElement(top_parallel, f"{{{P_NS}}}cTn", id="3", fill="hold")
    top_conditions = etree.SubElement(top_ctn, f"{{{P_NS}}}stCondLst")
    etree.SubElement(top_conditions, f"{{{P_NS}}}cond", delay="indefinite")
    top_children = etree.SubElement(top_ctn, f"{{{P_NS}}}childTnLst")
    behavior_parallel = etree.SubElement(top_children, f"{{{P_NS}}}par")
    behavior_ctn = etree.SubElement(behavior_parallel, f"{{{P_NS}}}cTn", id="4", fill="hold")
    behavior_conditions = etree.SubElement(behavior_ctn, f"{{{P_NS}}}stCondLst")
    etree.SubElement(behavior_conditions, f"{{{P_NS}}}cond", delay="0")
    behavior_children = etree.SubElement(behavior_ctn, f"{{{P_NS}}}childTnLst")
    effect_parallel = etree.SubElement(behavior_children, f"{{{P_NS}}}par")
    effect_ctn = etree.SubElement(
        effect_parallel,
        f"{{{P_NS}}}cTn",
        id="5",
        presetID=str(preset_id),
        presetClass="entr",
        presetSubtype="0",
        fill="hold",
        nodeType="clickEffect",
    )
    effect_conditions = etree.SubElement(effect_ctn, f"{{{P_NS}}}stCondLst")
    etree.SubElement(effect_conditions, f"{{{P_NS}}}cond", delay="0")
    animation_behavior(effect_ctn, 6, str(shape_id), effect)

    previous = etree.SubElement(sequence, f"{{{P_NS}}}prevCondLst")
    previous_condition = etree.SubElement(previous, f"{{{P_NS}}}cond", evt="onPrev", delay="0")
    previous_target = etree.SubElement(previous_condition, f"{{{P_NS}}}tgtEl")
    etree.SubElement(previous_target, f"{{{P_NS}}}sldTgt")
    following = etree.SubElement(sequence, f"{{{P_NS}}}nextCondLst")
    following_condition = etree.SubElement(following, f"{{{P_NS}}}cond", evt="onNext", delay="0")
    following_target = etree.SubElement(following_condition, f"{{{P_NS}}}tgtEl")
    etree.SubElement(following_target, f"{{{P_NS}}}sldTgt")
    return timing


def build_video_autoplay_timing(shape_id: int) -> etree._Element:
    """构造进入第 5 页后从 0 秒自动播放真实嵌入视频的时间线。"""
    timing = etree.Element(f"{{{P_NS}}}timing", nsmap=NSMAP)
    time_nodes = etree.SubElement(timing, f"{{{P_NS}}}tnLst")
    root_parallel = etree.SubElement(time_nodes, f"{{{P_NS}}}par")
    root_ctn = etree.SubElement(
        root_parallel,
        f"{{{P_NS}}}cTn",
        id="1",
        dur="indefinite",
        restart="never",
        nodeType="tmRoot",
    )
    root_children = etree.SubElement(root_ctn, f"{{{P_NS}}}childTnLst")

    sequence = etree.SubElement(root_children, f"{{{P_NS}}}seq", concurrent="1", nextAc="seek")
    main_ctn = etree.SubElement(
        sequence,
        f"{{{P_NS}}}cTn",
        id="2",
        dur="indefinite",
        nodeType="mainSeq",
    )
    main_children = etree.SubElement(main_ctn, f"{{{P_NS}}}childTnLst")
    top_parallel = etree.SubElement(main_children, f"{{{P_NS}}}par")
    top_ctn = etree.SubElement(top_parallel, f"{{{P_NS}}}cTn", id="3", fill="hold")
    top_conditions = etree.SubElement(top_ctn, f"{{{P_NS}}}stCondLst")
    etree.SubElement(top_conditions, f"{{{P_NS}}}cond", delay="0")
    top_children = etree.SubElement(top_ctn, f"{{{P_NS}}}childTnLst")
    inner_parallel = etree.SubElement(top_children, f"{{{P_NS}}}par")
    inner_ctn = etree.SubElement(inner_parallel, f"{{{P_NS}}}cTn", id="4", fill="hold")
    inner_conditions = etree.SubElement(inner_ctn, f"{{{P_NS}}}stCondLst")
    etree.SubElement(inner_conditions, f"{{{P_NS}}}cond", delay="0")
    inner_children = etree.SubElement(inner_ctn, f"{{{P_NS}}}childTnLst")
    command_parallel = etree.SubElement(inner_children, f"{{{P_NS}}}par")
    command_ctn = etree.SubElement(
        command_parallel,
        f"{{{P_NS}}}cTn",
        id="5",
        presetID="1",
        presetClass="mediacall",
        presetSubtype="0",
        fill="hold",
        nodeType="withEffect",
    )
    command_conditions = etree.SubElement(command_ctn, f"{{{P_NS}}}stCondLst")
    etree.SubElement(command_conditions, f"{{{P_NS}}}cond", delay="0")
    command_children = etree.SubElement(command_ctn, f"{{{P_NS}}}childTnLst")
    command = etree.SubElement(command_children, f"{{{P_NS}}}cmd", type="call", cmd="playFrom(0.0)")
    behavior = etree.SubElement(command, f"{{{P_NS}}}cBhvr")
    etree.SubElement(behavior, f"{{{P_NS}}}cTn", id="6", dur="10000", fill="hold")
    target = etree.SubElement(behavior, f"{{{P_NS}}}tgtEl")
    etree.SubElement(target, f"{{{P_NS}}}spTgt", spid=str(shape_id))

    previous = etree.SubElement(sequence, f"{{{P_NS}}}prevCondLst")
    previous_condition = etree.SubElement(previous, f"{{{P_NS}}}cond", evt="onPrev", delay="0")
    previous_target = etree.SubElement(previous_condition, f"{{{P_NS}}}tgtEl")
    etree.SubElement(previous_target, f"{{{P_NS}}}sldTgt")
    following = etree.SubElement(sequence, f"{{{P_NS}}}nextCondLst")
    following_condition = etree.SubElement(following, f"{{{P_NS}}}cond", evt="onNext", delay="0")
    following_target = etree.SubElement(following_condition, f"{{{P_NS}}}tgtEl")
    etree.SubElement(following_target, f"{{{P_NS}}}sldTgt")

    video = etree.SubElement(root_children, f"{{{P_NS}}}video", fullScrn="0")
    media = etree.SubElement(video, f"{{{P_NS}}}cMediaNode", vol="80000")
    media_ctn = etree.SubElement(media, f"{{{P_NS}}}cTn", id="7", fill="hold", display="1")
    media_conditions = etree.SubElement(media_ctn, f"{{{P_NS}}}stCondLst")
    etree.SubElement(media_conditions, f"{{{P_NS}}}cond", delay="0")
    media_target = etree.SubElement(media, f"{{{P_NS}}}tgtEl")
    etree.SubElement(media_target, f"{{{P_NS}}}spTgt", spid=str(shape_id))
    return timing


def set_timing(slide, timing: etree._Element) -> None:
    """以本页专用时间线替换原有动画。"""
    root = slide._element
    old = root.find(qn("p:timing"))
    if old is not None:
        root.remove(old)
    extension_list = root.find(qn("p:extLst"))
    if extension_list is None:
        root.append(timing)
    else:
        root.insert(root.index(extension_list), timing)


def first_picture(slide):
    pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    if not pictures:
        raise ValueError("目标页面没有可设置动画的图片")
    return pictures[0]


def reorder_slides(prs: Presentation) -> None:
    """把原第 6 页目录移到第 2 页，其余页面顺延且不改相对顺序。"""
    slide_ids = prs.slides._sldIdLst
    directory_id = slide_ids[5]
    slide_ids.remove(directory_id)
    slide_ids.insert(1, directory_id)


def apply_header_updates(slides) -> None:
    """逐页修改右上角章节标识。"""
    chapter_map = {
        3: "01 背景",
        4: "01 背景",
        5: "01 背景",
        6: "02 目标",
        7: "03 需求",
        8: "03 需求",
        9: "04 目标",
        10: "04 目标",
        11: "05 验证",
        12: "05 验证",
        13: "06 方案",
        14: "06 方案",
        15: "07 实施",
        16: "07 实施",
        17: "07 实施",
        18: "07 实施",
        19: "08 验证",
        20: "08 验证",
        21: "09 巩固",
    }
    for page_no, text in chapter_map.items():
        slide = slides[page_no - 1]
        candidates = [
            shape
            for shape in slide.shapes
            if shape.has_text_frame
            and shape.left >= Inches(10)
            and any(token in shape_text(shape) for token in ("背景", "目标", "需求", "验证", "方案", "实施", "巩固"))
        ]
        if len(candidates) != 1:
            raise ValueError(f"第 {page_no} 页章节标识命中 {len(candidates)} 个")
        shape = candidates[0]
        replace_text(shape, text, size=20, color=TEAL, bold=False, align=PP_ALIGN.CENTER)
        shape.top = Cm(1.05)
        shape.height = Cm(0.75)


def apply_task(input_dir: Path, output_path: Path) -> None:
    """按原始 23 页备注逐页完成真实任务。"""
    validate_inputs(input_dir)
    prs = Presentation(input_dir / SOURCE_NAME)
    if len(prs.slides) != 23:
        raise ValueError(f"真实输入应为 23 页，实际为 {len(prs.slides)} 页")
    reorder_slides(prs)
    slides = list(prs.slides)
    apply_header_updates(slides)

    # P2：目录标题及六个板块。
    directory = slides[1]
    title_shape = find_text_shape(directory, "目录", exact=True)
    replace_text(title_shape, "目录", size=40, color=WHITE, bold=True, align=PP_ALIGN.LEFT)
    title_shape.height = Inches(0.68)
    title_shape.text_frame.margin_top = 0
    title_shape.text_frame.margin_bottom = 0
    contents_shape = find_text_shape(directory, "CONTENTS", exact=True)
    replace_text(
        contents_shape,
        "CONTENTS",
        size=10,
        color=RGBColor(0x99, 0xF6, 0xE4),
        font="Arial",
        bold=True,
        align=PP_ALIGN.LEFT,
    )
    section_titles = [
        "项目背景与需求",
        "目标设定与可行性",
        "方案设计与优选",
        "实施过程",
        "效果验证",
        "标准化与展望",
    ]
    for index, text in enumerate(section_titles, 1):
        shape = find_text_shape(directory, text, exact=True)
        replace_text(shape, text, size=30, color=WHITE, bold=False, align=PP_ALIGN.LEFT)
        shape.height = Inches(0.50)
        shape.text_frame.margin_top = 0
        shape.text_frame.margin_bottom = 0
        number_shape = find_text_shape(directory, f"{index:02d}", exact=True)
        replace_text(number_shape, f"{index:02d}", size=30, color=WHITE, bold=False, align=PP_ALIGN.CENTER)
        number_shape.height = Inches(0.50)
        number_shape.text_frame.margin_top = 0
        number_shape.text_frame.margin_bottom = 0

    # P4：移除旧单图，并按原始文件顺序插入三张真实图片。
    slide4 = slides[3]
    for picture in [shape for shape in slide4.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]:
        remove_shape(picture)
    picture_positions = [
        (0.65, 1.48),
        (3.45, 1.54),
        (2.05, 4.02),
    ]
    for name, (left, top) in zip(IMAGE_NAMES, picture_positions):
        picture = slide4.shapes.add_picture(
            str(input_dir / name),
            Inches(left),
            Inches(top),
            width=Inches(2.50),
            height=Inches(1.875),
        )
        picture.line.color.rgb = LIGHT_BLUE
        picture.line.width = Pt(1.0)

    # P5：表格与真实嵌入视频自动播放。
    slide5 = slides[4]
    table5 = next(shape for shape in slide5.shapes if shape.has_table)
    format_table(table5, size=10, color=DEEP_BLUE)
    videos = [shape for shape in slide5.shapes if shape.shape_type == MSO_SHAPE_TYPE.MEDIA]
    if len(videos) != 1:
        raise ValueError(f"第 5 页视频对象数量应为 1，实际为 {len(videos)}")
    set_timing(slide5, build_video_autoplay_timing(videos[0].shape_id))

    # P6：三个目标关键词。
    slide6 = slides[5]
    keyword_colors = {
        "分离": TEAL,
        "观察": RGBColor(0x1D, 0x4E, 0xD8),
        "维护": RGBColor(0xF9, 0x73, 0x16),
    }
    for text, color in keyword_colors.items():
        replace_text(
            find_text_shape(slide6, text, exact=True),
            text,
            size=15,
            color=color,
            bold=True,
            align=PP_ALIGN.LEFT,
        )

    # P7：需求表格。
    format_table(next(shape for shape in slides[6].shapes if shape.has_table), size=10, color=DEEP_BLUE)

    # P8：三张原图统一为 8cm × 5cm。
    pictures8 = [shape for shape in slides[7].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    if len(pictures8) != 3:
        raise ValueError(f"第 8 页图片数量应为 3，实际为 {len(pictures8)}")
    for picture, left in zip(pictures8, (0.83, 5.09, 9.18)):
        picture.left = Inches(left)
        picture.top = Inches(3.43)
        picture.width = Cm(8)
        picture.height = Cm(5)

    # P9：把目标说明扩充到可公开交流的完整句子。
    slide9 = slides[8]
    target9 = find_text_shape(slide9, "目标设定采用区间化与功能化表达")
    expanded9 = "目标设定口径采用区间化与功能化表达，既保留可量化的设计边界，又避免绑定真实生产线专属数据，便于公开交流、横向比较与后续复核。"
    replace_text(target9, expanded9, size=20, color=TEAL, bold=False, align=PP_ALIGN.CENTER, line_spacing=0.95)
    target9.top = Inches(5.52)
    target9.height = Inches(1.02)
    target9.text_frame.margin_top = Inches(0.02)
    target9.text_frame.margin_bottom = Inches(0.02)

    # P10：主图片“出现”动画。
    picture10 = first_picture(slides[9])
    set_timing(slides[9], build_single_effect_timing(picture10.shape_id, preset_id=1, effect=None))

    # P11：扩充验证结论，并为主图片添加轮子动画。
    slide11 = slides[10]
    conclusion11 = find_text_shape(slide11, "单层筛面已经能够形成稳定分路")
    expanded11 = "单层筛面已经能够形成稳定分路，后续将围绕不同物料粒径持续优化孔型，联合校正筛面角度与输送节拍，并完善密封边界优化和清理验证，以降低粉料外逸与局部堆积风险。"
    replace_text(conclusion11, expanded11, size=13, color=BLACK, bold=False, align=PP_ALIGN.LEFT, line_spacing=0.95)
    conclusion11.top = Inches(4.70)
    conclusion11.height = Inches(0.98)
    set_timing(
        slide11,
        build_single_effect_timing(first_picture(slide11).shape_id, preset_id=21, effect="wheel"),
    )

    # P12：试验说明。
    slide12 = slides[11]
    statement12 = find_text_shape(slide12, "本试验仅用于工艺趋势判断")
    replace_text(
        statement12,
        "本试验仅用于工艺趋势判断，不作为真实生产环境的燃烧测试记录。",
        size=20,
        color=SLATE,
        bold=False,
        align=PP_ALIGN.CENTER,
    )
    statement12.top = Inches(6.28)
    statement12.height = Inches(0.56)

    # P13：将设计原则人工扩展为五条，保留原版布局和通用设备图。
    slide13 = slides[12]
    design = find_text_shape(slide13, "设计原则")
    replace_text(
        design,
        "设计原则\n1. 不干扰主线控制\n2. 不占用关键通道\n3. 保留分路可视性\n4. 便于拆装与清洁\n5. 兼顾密封与扩展",
        size=12.5,
        color=DEEP_BLUE,
        bold=False,
        align=PP_ALIGN.LEFT,
        line_spacing=0.86,
    )
    design.left = Inches(6.32)
    design.top = Inches(4.84)
    design.width = Inches(5.58)
    design.height = Inches(1.58)
    design.text_frame.margin_top = Inches(0.03)
    design.text_frame.margin_bottom = Inches(0.03)
    for shape in slide13.shapes:
        if shape is design or shape.shape_type == MSO_SHAPE_TYPE.PICTURE or shape.has_table:
            continue
        if (
            not shape_text(shape)
            and abs(shape.left - Inches(6.10)) < Inches(0.05)
            and shape.top >= Inches(5.0)
        ):
            shape.top = Inches(4.76)
            shape.height = Inches(1.78)
            break

    # P14：研发指标表格和底部公开口径说明。
    slide14 = slides[13]
    table14 = next(shape for shape in slide14.shapes if shape.has_table)
    table14.top = Inches(1.20)
    table14.height = Inches(4.82)
    for row in table14.table.rows:
        row.height = Inches(4.82 / 6)
    format_table(table14, size=20, color=DEEP_BLUE)
    footer14 = find_text_shape(slide14, "指标已经转为公开交流口径")
    replace_text(
        footer14,
        "指标已经转为公开交流口径，不包含原采购金额、供应渠道与真实硬件型号。",
        size=20,
        color=SLATE,
        bold=False,
        align=PP_ALIGN.CENTER,
    )
    footer14.top = Inches(6.23)
    footer14.height = Inches(0.58)

    # P15：责任角色精确替换，表格保持原生可编辑。
    slide15 = slides[14]
    table15 = next(shape for shape in slide15.shapes if shape.has_table)
    table15.table.cell(4, 3).text = "项目协调负责人：沈亦然"
    format_table(table15, size=10, color=DEEP_BLUE)

    # P16：三条实施说明。
    slide16 = slides[15]
    for text in (
        "长孔筛面兼顾分离效率与堵塞风险",
        "让物料在自重与振动下自然前移",
        "减少粉料外逸和箱体边缘堆积",
    ):
        shape = find_text_shape(slide16, text, exact=True)
        replace_text(shape, text, size=15, color=DEEP_BLUE, bold=False, align=PP_ALIGN.LEFT)
        shape.height = Inches(0.55)

    # P17：表格与主图渐变动画。
    slide17 = slides[16]
    format_table(next(shape for shape in slide17.shapes if shape.has_table), size=10, color=DEEP_BLUE)
    set_timing(
        slide17,
        build_single_effect_timing(first_picture(slide17).shape_id, preset_id=10, effect="fade"),
    )

    # P18：主图溶解动画。
    slide18 = slides[17]
    set_timing(
        slide18,
        build_single_effect_timing(first_picture(slide18).shape_id, preset_id=1, effect="dissolve"),
    )

    # P19：用真实工作簿的 A1:D6 数据替换表格。
    replace_excel_table(slides[18], input_dir / WORKBOOK_NAME)

    # P20：风险链条两个圆形颜色。
    slide20 = slides[19]
    set_shape_fill_and_line(find_circle_above(slide20, "燃料侧"), RGBColor(0xDC, 0x26, 0x26))
    set_shape_fill_and_line(find_circle_above(slide20, "发现侧"), RGBColor(0xFA, 0xCC, 0x15))

    # P21：巩固措施说明。
    slide21 = slides[20]
    next_step = find_text_shape(slide21, "下一步：根据不同线体空间差异")
    replace_text(
        next_step,
        "下一步：根据不同线体空间差异，形成可调整的安装包和巡检说明卡。",
        size=20,
        color=NAVY,
        bold=False,
        align=PP_ALIGN.LEFT,
    )
    next_step.top = Inches(5.42)
    next_step.height = Inches(0.58)

    # P22：总结页标题和后续方向。
    slide22 = slides[21]
    summary_title = find_text_shape(slide22, "总结与展望", exact=True)
    replace_text(summary_title, "总结与展望", size=40, color=WHITE, bold=True, align=PP_ALIGN.LEFT)
    summary_title.height = Inches(0.70)
    summary_footer = find_text_shape(slide22, "后续方向：结合图像识别与清理记录")
    replace_text(
        summary_footer,
        "后续方向：结合图像识别与清理记录，实现异常趋势提醒与点检数据沉淀。",
        size=20,
        color=WHITE,
        bold=False,
        align=PP_ALIGN.LEFT,
    )
    summary_footer.height = Inches(0.55)

    # P23：感谢页标题。
    thanks = find_text_shape(slides[22], "感谢聆听", exact=True)
    replace_text(thanks, "感谢聆听", size=40, color=WHITE, bold=True, align=PP_ALIGN.LEFT)
    thanks.height = Inches(0.72)

    # 除 P19 外，所有表格按原始 Rubric 要求统一左对齐。
    for page_no, slide in enumerate(slides, 1):
        if page_no == 19:
            continue
        for shape in slide.shapes:
            if not shape.has_table:
                continue
            for row in shape.table.rows:
                for cell in row.cells:
                    for paragraph in cell.text_frame.paragraphs:
                        paragraph.alignment = PP_ALIGN.LEFT

    prs.core_properties.title = "高速成型线碎屑分离与状态监测装置改进汇报"
    prs.core_properties.subject = "OmegaUse-OfficeVal officeval_041 逐页备注修改参考交付"
    prs.core_properties.author = "Cowork Evolution Benchmark"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    arguments = parser.parse_args()
    apply_task(arguments.input_dir, arguments.output_dir / OUTPUT_NAME)


if __name__ == "__main__":
    main()
