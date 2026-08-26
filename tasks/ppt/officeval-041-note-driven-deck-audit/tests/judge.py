#!/usr/bin/env python3
"""officeval_041 逐页备注工程演示文稿任务的 Harbor Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import upstream_verifier


OUTPUT_NAME = "高速成型线碎屑分离与状态监测装置_修改完成版.pptx"
TASK_ID = "cowork-evo/officeval-041-note-driven-deck-audit"
RUBRIC_VERSION = "omegause-officeval-041-adaptation-v0.1"
JUDGE_VERSION = "officeval-041-adaptation-v0.1"
PASS_THRESHOLD = 0.8
VIDEO_SHA256 = "9ff085308ac735e9a97555084a4c2985630591f8aad559b0ee7ebb89a0a55601"
IMAGE_HASHES = [
    "cfdaf30c5c63be8d72d2e3a67e155bbdd3c46913a5784ca5b3bf61852c150150",
    "57de004c312ce1a7dd6eaee93bcdf876a8f1d9ee0aecf875dcf16571463546a8",
    "88573eede478416ac4226dedc802d804fc29a5592e022f214464bf8915638b51",
]

EXPECTED_TITLES = [
    ["高速成型线碎屑分离与状态监测装置改进汇报"],
    ["目录", "CONTENTS"],
    ["项目定位"],
    ["风险机制", "三个要素相互叠加"],
    ["改造前", "回收路径长", "观察节点少"],
    ["改造目标", "把风险链拆开"],
    ["需求转换", "从现场声音到关键需求"],
    ["创新借鉴", "分割与预先感知"],
    ["目标设定", "用可量化指标约束设计"],
    ["可行性论证", "从工艺", "技术", "资源三方面评估"],
    ["模拟试验一", "筛分路径验证"],
    ["模拟试验二", "粉料含量与异常蔓延趋势"],
    ["总体方案", "物理分路", "在线状态感知"],
    ["研发指标", "功能", "性能与兼容性"],
    ["制定对策", "任务", "目标与交付物"],
    ["对策实施一", "筛分系统设计与加工"],
    ["对策实施二", "可开合观察窗制作"],
    ["对策实施三", "监测画面接入操作位"],
    ["整体试运行", "连续班次结果"],
    ["效果验证", "风险链条被削弱"],
    ["标准化与巩固措施"],
    ["总结与展望"],
    ["感谢聆听"],
]

CRITERION_SPECS = {
    "R001": ("交付指定名称的可打开 23 页 PPTX", 2, "hurdle"),
    "R002": ("目录移到第 2 页且其余页面顺序正确", 5, "hurdle"),
    "R003": ("目录与章节标识格式正确", 5, "positive"),
    "R004": ("P4/P8 图片来源、尺寸、边框和遮挡正确", 4, "positive"),
    "R005": ("逐页备注的文本、表格、字体和版式修改完整", 6, "positive"),
    "R006": ("P19 与真实 Excel A1:D6 逐格一致", 5, "hurdle"),
    "R007": ("真实视频自动播放且四页图片动画正确", 5, "positive"),
    "R008": ("颜色、末页和全部负向约束正确", 4, "positive"),
    "R009": ("修正缺陷后上游规则满分", 5, "positive"),
    "R010": ("LibreOffice 实际渲染 23 页", 3, "positive"),
    "R011": ("文本表格可编辑且视频字节保持", 3, "positive"),
    "R012": ("关键内容、顺序或 Excel 不一致", -10, "penalty"),
    "R013": ("视频或指定动画缺失", -8, "penalty"),
    "R014": ("文件损坏、页数错误或多数页面压平", -6, "penalty"),
}


def criterion(identifier: str, passed: bool, evidence: str = "", score: float | None = None) -> dict:
    description, weight, kind = CRITERION_SPECS[identifier]
    if score is None:
        score = 1.0 if passed else 0.0
    return {
        "criterion_id": identifier,
        "description": description,
        "weight": weight,
        "type": kind,
        "passed": bool(passed),
        "score": max(0.0, min(1.0, float(score))),
        "evidence": evidence,
        "method": "deterministic",
    }


def normalized(text: str) -> str:
    return re.sub(r"\s+", "", text or "")


def slide_text(slide) -> str:
    return "\n".join(
        shape.text_frame.text
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )


def check_order(prs: Presentation) -> tuple[bool, str]:
    if len(prs.slides) != len(EXPECTED_TITLES):
        return False, f"页数={len(prs.slides)}"
    errors = []
    for page_no, (slide, expected) in enumerate(zip(prs.slides, EXPECTED_TITLES), 1):
        text = normalized(slide_text(slide))
        missing = [item for item in expected if normalized(item) not in text]
        if missing:
            errors.append(f"P{page_no} 缺少 {missing}")
    return not errors, "23 页标题与目录顺序正确" if not errors else "；".join(errors[:10])


def source_workbook_values() -> list[list[str]]:
    task_root = Path(__file__).resolve().parents[1]
    path = task_root / "data" / "input_files" / "对比分析表.xlsx"
    workbook = load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook["Sheet1"]
    values = [
        ["" if worksheet.cell(row=row, column=column).value is None else str(worksheet.cell(row=row, column=column).value)
         for column in range(1, 5)]
        for row in range(1, 7)
    ]
    workbook.close()
    return values


def check_workbook_table(prs: Presentation) -> tuple[bool, str]:
    tables = [shape for shape in prs.slides[18].shapes if shape.has_table]
    if len(tables) != 1:
        return False, f"P19 表格数量={len(tables)}"
    table_shape = tables[0]
    actual_rows = [
        [cell.text for cell in list(row.cells)[:4]]
        for row in list(table_shape.table.rows)[:6]
    ]
    expected = source_workbook_values()
    if len(actual_rows) != 6 or any(len(row) != 4 for row in actual_rows):
        width_cm = table_shape.width / 360000
        height_cm = table_shape.height / 360000
        return False, (
            f"P19 表格结构={len(actual_rows)}行×"
            f"{max((len(row) for row in actual_rows), default=0)}列，"
            f"要求6行×4列；尺寸={width_cm:.2f}×{height_cm:.2f}cm"
        )
    actual = actual_rows
    mismatches = []
    for row in range(6):
        for column in range(4):
            if actual[row][column] != expected[row][column]:
                mismatches.append(
                    f"{chr(65 + column)}{row + 1}: {actual[row][column]!r} != {expected[row][column]!r}"
                )
    width_cm = table_shape.width / 360000
    height_cm = table_shape.height / 360000
    size_ok = abs(width_cm - 24.07) <= 0.2 and abs(height_cm - 10.59) <= 0.2
    ok = not mismatches and size_ok
    evidence = (
        f"24/24 单元格与真实 Excel 一致，尺寸={width_cm:.2f}×{height_cm:.2f}cm"
        if ok
        else f"差异={mismatches[:8]}，尺寸={width_cm:.2f}×{height_cm:.2f}cm"
    )
    return ok, evidence


def run_upstream(artifact_dir: Path) -> dict:
    try:
        return upstream_verifier.evaluate(str(artifact_dir))
    except Exception as exc:
        return {
            "status": "error",
            "error": f"{type(exc).__name__}: {exc}",
            "dim1_pass": False,
            "dim2_items": [],
            "total_score": 0,
            "max_score": 66,
        }


def upstream_rule_map(result: dict) -> dict[str, dict]:
    return {item.get("rule", ""): item for item in result.get("dim2_items", [])}


def rules_passed(rule_map: dict[str, dict], fragments: list[str]) -> tuple[bool, str]:
    evidence = []
    ok = True
    for fragment in fragments:
        matches = [item for rule, item in rule_map.items() if fragment in rule]
        passed = len(matches) == 1 and bool(matches[0].get("hit"))
        evidence.append(f"{fragment}={'通过' if passed else '失败'}")
        ok = ok and passed
    return ok, "；".join(evidence)


def check_images(prs: Presentation) -> tuple[bool, str]:
    pictures4 = [shape for shape in prs.slides[3].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    hashes4 = [hashlib.sha256(shape.image.blob).hexdigest() for shape in pictures4]
    source_ok = hashes4 == IMAGE_HASHES
    ratio_border_ok = True
    for shape in pictures4:
        ratio = shape.width / shape.height if shape.height else 0
        xml = shape._element.xml.upper()
        ratio_border_ok = ratio_border_ok and 1.13 <= ratio <= 1.53 and "93C5FD" in xml

    pictures8 = [shape for shape in prs.slides[7].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    size8_ok = len(pictures8) == 3 and all(
        abs(shape.width / 360000 - 8.0) <= 0.2
        and abs(shape.height / 360000 - 5.0) <= 0.2
        for shape in pictures8
    )
    ok = len(pictures4) == 3 and source_ok and ratio_border_ok and size8_ok
    return ok, f"P4图片哈希={hashes4}；P4比例边框={ratio_border_ok}；P8尺寸={size8_ok}"


def embedded_video_hash(path: Path) -> str | None:
    try:
        with ZipFile(path) as archive:
            videos = [name for name in archive.namelist() if name.startswith("ppt/media/") and name.lower().endswith(".mp4")]
            if len(videos) != 1:
                return None
            return hashlib.sha256(archive.read(videos[0])).hexdigest()
    except Exception:
        return None


def check_media_and_animations(prs: Presentation, path: Path) -> tuple[bool, str]:
    video_hash = embedded_video_hash(path)
    xml5 = prs.slides[4]._element.xml
    video_ok = (
        video_hash == VIDEO_SHA256
        and 'nodeType="mainSeq"' in xml5
        and 'cmd="playFrom(0.0)"' in xml5
        and 'delay="0"' in xml5
    )
    expected = {
        10: ['presetID="1"', 'presetClass="entr"'],
        11: ['presetID="21"', 'filter="wheel"'],
        17: ['presetID="10"', 'filter="fade"'],
        18: ['filter="dissolve"'],
    }
    page_results = {}
    for page_no, tokens in expected.items():
        xml = prs.slides[page_no - 1]._element.xml
        picture_ids = {
            str(shape.shape_id)
            for shape in prs.slides[page_no - 1].shapes
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        }
        targets = set(re.findall(r'<p:spTgt\s+spid="(\d+)"', xml))
        page_results[page_no] = all(token in xml for token in tokens) and bool(picture_ids & targets)
    ok = video_ok and all(page_results.values())
    return ok, f"视频哈希={video_hash}，自动播放={video_ok}，动画页={page_results}"


def check_editability(prs: Presentation, path: Path) -> tuple[bool, str, int]:
    editable_text = 0
    editable_tables = 0
    flattened_pages = 0
    for page_no, slide in enumerate(prs.slides, 1):
        page_text_shapes = 0
        oversized_pictures = 0
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                editable_text += 1
                page_text_shapes += 1
            if getattr(shape, "has_table", False):
                editable_tables += 1
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                if shape.width >= prs.slide_width * 0.9 and shape.height >= prs.slide_height * 0.9:
                    oversized_pictures += 1
        if oversized_pictures and page_text_shapes == 0:
            flattened_pages += 1
    ok = editable_text >= 120 and editable_tables >= 8 and flattened_pages == 0 and embedded_video_hash(path) == VIDEO_SHA256
    return ok, f"可编辑文本对象={editable_text}，表格={editable_tables}，压平页={flattened_pages}", flattened_pages


def render_check(path: Path) -> tuple[bool, str]:
    libreoffice = shutil.which("libreoffice")
    pdfinfo = shutil.which("pdfinfo")
    pdftotext = shutil.which("pdftotext")
    if not libreoffice or not pdfinfo or not pdftotext:
        return False, "缺少 libreoffice、pdfinfo 或 pdftotext"
    with tempfile.TemporaryDirectory(prefix="cowork-evo-041-") as temp_dir_text:
        temp_dir = Path(temp_dir_text)
        profile = temp_dir / "profile"
        profile.mkdir()
        environment = os.environ.copy()
        for key in ("http_proxy", "https_proxy", "HTTP_PROXY", "HTTPS_PROXY", "ALL_PROXY", "all_proxy"):
            environment.pop(key, None)
        converted = subprocess.run(
            [
                libreoffice,
                "--headless",
                f"-env:UserInstallation=file://{profile}",
                "--convert-to",
                "pdf",
                "--outdir",
                str(temp_dir),
                str(path),
            ],
            capture_output=True,
            text=True,
            timeout=300,
            check=False,
            env=environment,
        )
        pdf_path = temp_dir / f"{path.stem}.pdf"
        if converted.returncode != 0 or not pdf_path.is_file():
            return False, (converted.stderr or converted.stdout or "转换失败")[-1000:]
        info = subprocess.run([pdfinfo, str(pdf_path)], capture_output=True, text=True, timeout=60, check=False)
        extracted = subprocess.run([pdftotext, "-layout", str(pdf_path), "-"], capture_output=True, text=True, timeout=120, check=False)
        match = re.search(r"^Pages:\s+(\d+)", info.stdout, re.M)
        pages = int(match.group(1)) if match else 0
        text = normalized(extracted.stdout)
        anchors = [
            "高速成型线碎屑分离与状态监测装置改进汇报",
            "目录",
            "总体方案：物理分路+在线状态感知",
            "整体试运行：连续班次结果",
            "感谢聆听",
        ]
        missing = [anchor for anchor in anchors if normalized(anchor) not in text]
    ok = pages == 23 and extracted.returncode == 0 and not missing
    return ok, f"渲染页数={pages}，缺失关键文字={missing}"


def build_result(results: list[dict], artifact_path: Path) -> dict:
    present = {item["criterion_id"] for item in results}
    for identifier in CRITERION_SPECS:
        if identifier not in present:
            results.append(criterion(identifier, False, "因交付文件不存在、损坏或门槛失败，本判据未执行"))
    results.sort(key=lambda item: item["criterion_id"])
    positive_total = sum(item["weight"] for item in results if item["weight"] > 0)
    earned = sum(item["weight"] * item["score"] for item in results if item["weight"] > 0)
    penalties = sum(item["weight"] for item in results if item["weight"] < 0 and item["passed"])
    reward = max(0.0, min(1.0, (earned + penalties) / positive_total))
    hurdle_failures = [item["criterion_id"] for item in results if item["type"] == "hurdle" and not item["passed"]]
    if hurdle_failures:
        reward = 0.0
    digest = hashlib.sha256(artifact_path.read_bytes()).hexdigest() if artifact_path.is_file() else hashlib.sha256(b"").hexdigest()
    normalized_results = []
    for item in results:
        normalized_score = (0.0 if item["passed"] else 1.0) if item["type"] == "penalty" else item["score"]
        normalized_results.append({
            "criterion_id": item["criterion_id"],
            "score": normalized_score,
            "evidence": item["evidence"],
            "method": item["method"],
            "raw": {
                "description": item["description"],
                "weight": item["weight"],
                "type": item["type"],
                "passed": item["passed"],
            },
        })
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [digest],
        "criterion_results": normalized_results,
        "reward": round(reward, 6),
        "passed": reward >= PASS_THRESHOLD,
        "hurdle_failures": hurdle_failures,
    }


def evaluate(artifact_dir: Path) -> dict:
    output_path = artifact_dir / OUTPUT_NAME
    results = []
    if not output_path.is_file():
        results.append(criterion("R001", False, f"缺少 {OUTPUT_NAME}"))
        return build_result(results, output_path)
    try:
        prs = Presentation(output_path)
    except Exception as exc:
        results.append(criterion("R001", False, f"文件无法打开：{exc}"))
        results.append(criterion("R014", True, "交付文件损坏"))
        return build_result(results, output_path)

    integrity_ok = len(prs.slides) == 23
    results.append(criterion("R001", integrity_ok, f"页数={len(prs.slides)}"))
    if not integrity_ok:
        results.append(criterion("R014", True, f"页数={len(prs.slides)}"))
        return build_result(results, output_path)

    order_ok, order_evidence = check_order(prs)
    results.append(criterion("R002", order_ok, order_evidence))
    workbook_ok, workbook_evidence = check_workbook_table(prs)
    results.append(criterion("R006", workbook_ok, workbook_evidence))

    upstream = run_upstream(artifact_dir)
    rule_map = upstream_rule_map(upstream)
    directory_ok, directory_evidence = rules_passed(
        rule_map,
        ["章节标识文字组", "目录页标题文字", "目录六个板块文字"],
    )
    results.append(criterion("R003", directory_ok, directory_evidence))

    images_ok, images_evidence = check_images(prs)
    upstream_images_ok, upstream_images_evidence = rules_passed(
        rule_map,
        ["第4页三张图片插入", "第8页三张图尺寸统一"],
    )
    results.append(criterion("R004", images_ok and upstream_images_ok, f"{images_evidence}；{upstream_images_evidence}"))

    note_fragments = [
        "第5页表格文字",
        "第6页目标关键词",
        "第7页需求表格",
        "第9页目标说明文字扩充",
        "第11页验证结论扩充",
        "第12页试验说明文字",
        "第13页设计原则",
        "第14页研发指标表格",
        "第15页'项目协调负责人",
        "第16页三句话",
        "第17页表格文字",
        "第19页Excel表格尺寸和表头",
        "第20页燃料侧红色",
        "第21页巩固措施说明",
        "第22页总结与展望文字",
        "第23页感谢页标题",
    ]
    notes_ok, notes_evidence = rules_passed(rule_map, note_fragments)
    results.append(criterion("R005", notes_ok, notes_evidence))

    media_ok, media_evidence = check_media_and_animations(prs, output_path)
    upstream_animation_ok, upstream_animation_evidence = rules_passed(
        rule_map,
        ["第5页视频自动播放", "第10页图片出现动画", "第11页图片轮子动画", "第17页图片渐变动画", "第18页溶解动画"],
    )
    results.append(criterion("R007", media_ok and upstream_animation_ok, f"{media_evidence}；{upstream_animation_evidence}"))

    deductions = [item for item in upstream.get("dim2_items", []) if item.get("max_delta", 0) < 0 and item.get("hit")]
    results.append(criterion("R008", not deductions, f"上游扣分项触发={[(item.get('rule'), item.get('delta')) for item in deductions]}"))

    positive_failures = [
        item.get("rule")
        for item in upstream.get("dim2_items", [])
        if item.get("max_delta", 0) > 0
        and not item.get("hit")
        and "第19页Excel表格具体数据" not in item.get("rule", "")
    ]
    corrected_full = (
        upstream.get("status") == "ok"
        and upstream.get("dim1_pass") is True
        and upstream.get("total_score") == 65
        and upstream.get("max_score") == 66
        and not positive_failures
        and not deductions
        and workbook_ok
    )
    results.append(
        criterion(
            "R009",
            corrected_full,
            f"上游={upstream.get('total_score')}/{upstream.get('max_score')}；非已知缺陷失败={positive_failures}；真实Excel修正={workbook_ok}",
        )
    )

    rendered, render_evidence = render_check(output_path)
    results.append(criterion("R010", rendered, render_evidence))
    editable, editable_evidence, flattened_pages = check_editability(prs, output_path)
    results.append(criterion("R011", editable, editable_evidence))

    critical_fidelity_failure = not (order_ok and workbook_ok and notes_ok)
    results.append(criterion("R012", critical_fidelity_failure, f"顺序={order_ok}，Excel={workbook_ok}，备注修改={notes_ok}"))
    results.append(criterion("R013", not (media_ok and upstream_animation_ok), media_evidence))
    severe_failure = not integrity_ok or flattened_pages >= 12
    results.append(criterion("R014", severe_failure, f"完整性={integrity_ok}，压平页={flattened_pages}"))
    return build_result(results, output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path)
    arguments = parser.parse_args()
    result = evaluate(arguments.artifact_dir)
    rendered = json.dumps(result, ensure_ascii=False, indent=2)
    print(rendered)
    if arguments.result:
        arguments.result.parent.mkdir(parents=True, exist_ok=True)
        arguments.result.write_text(rendered + "\n", encoding="utf-8")
    raise SystemExit(0 if result["passed"] else 1)


if __name__ == "__main__":
    main()
