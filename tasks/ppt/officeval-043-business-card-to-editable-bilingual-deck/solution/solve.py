#!/usr/bin/env python3
"""officeval_043 逐题解法：依据名片参考图重建两页可编辑 PPT。"""

from __future__ import annotations

import argparse
import hashlib
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Pt


INPUT_NAME = "image_001.png"
OUTPUT_NAME = "lanyu_business_card_editable.pptx"
INPUT_SHA256 = "8a504c07410fd9c6d092fd43dc8943391c92c34e9368cd9918e09d745b338124"

GREEN = "075B51"
DARK_GREEN = "064B43"
MINT = "DDE8DF"
ORANGE = "C87738"
PALE_ORANGE = "E8B27E"
INK = "152D2B"
GRAY = "EBEEE9"
WHITE = "FFFEF9"


def color(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def add_text(slide, text: str, x: float, y: float, w: float, h: float, *,
             size: float = 12, font: str = "Arial", fill: str = INK,
             bold: bool = False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.MIDDLE, italic: bool = False):
    box = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    box.name = f"text-{text[:24]}"
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Cm(0.03)
    frame.margin_right = Cm(0.03)
    frame.margin_top = Cm(0.02)
    frame.margin_bottom = Cm(0.02)
    frame.vertical_anchor = anchor
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color(fill)
    return box


def add_shape(slide, kind, x: float, y: float, w: float, h: float, *,
              fill: str = WHITE, line: str | None = None, width: float = 1.0,
              radius_name: str | None = None):
    shape = slide.shapes.add_shape(kind, Cm(x), Cm(y), Cm(w), Cm(h))
    shape.name = radius_name or f"shape-{len(slide.shapes)}"
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = color(line)
        shape.line.width = Pt(width)
    return shape


def add_line(slide, x1: float, y1: float, x2: float, y2: float, *,
             fill: str = GREEN, width: float = 1.0, dash: bool = False):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    line.line.color.rgb = color(fill)
    line.line.width = Pt(width)
    if dash:
        line.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return line


def add_leaf(slide, x: float, y: float, scale: float = 1.0, *, fill: str = GREEN):
    leaf = add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x, y, 0.34 * scale, 0.68 * scale, fill=fill)
    leaf.rotation = 35
    add_line(slide, x + 0.07 * scale, y + 0.52 * scale, x + 0.30 * scale, y + 0.12 * scale,
             fill=WHITE, width=0.7)
    return leaf


def add_logo(slide, x: float, y: float, scale: float = 1.0):
    # 原图左上角的双叶与日轮，均由可编辑形状组成。
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + 0.92 * scale, y, 0.42 * scale, 0.42 * scale,
              fill=ORANGE, radius_name="logo-sun")
    left = add_shape(slide, MSO_AUTO_SHAPE_TYPE.TEAR, x, y + 0.35 * scale, 1.0 * scale, 1.55 * scale,
                     fill=GREEN, radius_name="logo-left-leaf")
    left.rotation = 315
    right = add_shape(slide, MSO_AUTO_SHAPE_TYPE.TEAR, x + 0.72 * scale, y + 0.38 * scale,
                      1.0 * scale, 1.55 * scale, fill=DARK_GREEN, radius_name="logo-right-leaf")
    right.rotation = 45
    add_line(slide, x + 0.22 * scale, y + 1.55 * scale, x + 1.62 * scale, y + 1.55 * scale,
             fill=GREEN, width=1.7)


def add_leaf_divider(slide, x: float, y: float, total_w: float, *, line_color: str = ORANGE):
    center = x + total_w / 2
    add_line(slide, x, y, center - 0.35, y, fill=line_color, width=1.0)
    add_leaf(slide, center - 0.15, y - 0.23, 0.52, fill=ORANGE if line_color == ORANGE else GREEN)
    add_line(slide, center + 0.35, y, x + total_w, y, fill=line_color, width=1.0)


def add_contact_icon(slide, symbol: str, x: float, y: float):
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, 0.9, 0.9,
              fill=DARK_GREEN, radius_name=f"contact-{symbol}")
    add_text(slide, symbol, x + 0.06, y + 0.03, 0.78, 0.78, size=17,
             font="DejaVu Sans", fill=WHITE, bold=True, align=PP_ALIGN.CENTER)


def add_business_icon(slide, kind: str, x: float, y: float):
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x, y, 1.4, 1.4, fill=GRAY, radius_name=f"business-{kind}-bg")
    if kind == "drop":
        drop = add_shape(slide, MSO_AUTO_SHAPE_TYPE.TEAR, x + 0.45, y + 0.20, 0.58, 0.86,
                         fill="C7DED3", line=GREEN, width=1.2, radius_name="business-drop")
        drop.rotation = 0
        add_leaf(slide, x + 0.77, y + 0.72, 0.48)
    elif kind == "leaf":
        add_leaf(slide, x + 0.37, y + 0.22, 1.35)
        add_line(slide, x + 0.50, y + 1.08, x + 1.02, y + 0.38, fill=WHITE, width=1.0)
    else:
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.CUBE, x + 0.30, y + 0.30, 0.82, 0.82,
                  fill=GREEN, line=WHITE, width=1.0, radius_name="business-box")


def add_business_list(slide, x: float, y: float, *, english: bool):
    headings = (
        [("ECO COATINGS:", ["Water-based resins", "Functional additives", "Barrier coatings"]),
         ("GREEN MATERIALS:", ["Biodegradable films", "Paper-plastic composites"]),
         ("APPLICATION SOLUTIONS:", ["Industrial adhesives", "Packaging surface treatment"])]
        if english else
        [("环保涂层：", ["水性树脂", "功能助剂", "高阻隔涂层"]),
         ("绿色材料：", ["可降解薄膜", "纸塑复合材料"]),
         ("应用方案：", ["工业胶黏剂", "包装表面处理"])]
    )
    icon_kinds = ["drop", "leaf", "box"]
    cursor = y
    font = "Arial" if english else "Microsoft YaHei"
    for index, (heading, items) in enumerate(headings):
        add_business_icon(slide, icon_kinds[index], x, cursor + 0.05)
        add_text(slide, heading, x + 1.8, cursor, 5.1, 0.45,
                 size=11.5 if english else 13, font=font, fill=GREEN, bold=True)
        for row, item in enumerate(items):
            add_text(slide, item, x + 1.8, cursor + 0.48 + row * 0.43, 5.2, 0.42,
                     size=9.5 if english else 11, font=font, fill=INK)
        block_h = 1.95 if len(items) == 3 else 1.55
        cursor += block_h
        if index < 2:
            divider = add_line(slide, x, cursor - 0.15, x + 6.6, cursor - 0.15,
                               fill=ORANGE, width=0.7, dash=True)
            divider.name = f"business-divider-{index + 1}"


def add_common_frame(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color("F7F5EE")
    frame = add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                      4.23, 2.38, 25.4, 14.3, fill="FBFAF4", line="333333", width=1.1,
                      radius_name="business-card-frame")
    frame.adjustments[0] = 0.08


def add_qr(slide, qr_path: Path):
    picture = slide.shapes.add_picture(str(qr_path), Cm(25.25), Cm(3.05), Cm(3.25), Cm(3.25))
    picture.name = "qr-code-source-crop"


def add_slide_one(prs: Presentation, qr_path: Path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_common_frame(slide)
    add_logo(slide, 5.45, 3.28, 1.05)
    add_text(slide, "福建蓝屿新材料科技有限公司", 8.0, 3.35, 15.8, 0.82,
             size=28, font="Microsoft YaHei", fill=DARK_GREEN, bold=True)
    add_text(slide, "FUJIAN LANYU NOVA MATERIALS CO., LTD.", 8.1, 4.18, 13.3, 0.52,
             size=16, font="Arial", fill=ORANGE)
    add_leaf_divider(slide, 8.1, 4.93, 13.3)
    add_qr(slide, qr_path)

    add_text(slide, "沈知行", 6.8, 6.5, 7.6, 1.45, size=45, font="STKaiti",
             fill=DARK_GREEN, bold=True)
    add_leaf_divider(slide, 7.0, 8.55, 6.2, line_color=GREEN)
    add_contact_icon(slide, "☎", 6.7, 10.25)
    add_text(slide, "M.P.   139 4826 7153", 7.9, 10.25, 7.5, 0.9, size=18,
             font="Arial", fill=INK, bold=True)
    add_contact_icon(slide, "✉", 6.7, 11.45)
    add_text(slide, "z.xing@lanyunova.com", 7.9, 11.45, 7.5, 0.9, size=15,
             font="Arial", fill=INK)
    add_contact_icon(slide, "⌖", 6.7, 12.65)
    add_text(slide, "福建省厦门市海沧区云创大道88号\n蓝湾创新中心12层1208室", 7.9, 12.55, 8.0, 1.3,
             size=11, font="Microsoft YaHei", fill=INK, anchor=MSO_ANCHOR.TOP)
    add_line(slide, 17.05, 6.95, 17.05, 14.65, fill="B7C8BE", width=1.0)

    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 18.15, 6.65, 3.4, 0.72,
              fill=GREEN, radius_name="main-business-pill")
    add_text(slide, "主  营", 18.15, 6.66, 3.4, 0.68, size=15, font="Microsoft YaHei",
             fill=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_business_list(slide, 18.3, 7.65, english=False)

    add_shape(slide, MSO_AUTO_SHAPE_TYPE.WAVE, 4.28, 15.25, 25.28, 1.35,
              fill=GREEN, radius_name="bottom-green-wave")
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.WAVE, 4.28, 15.72, 10.2, 0.82,
              fill=PALE_ORANGE, radius_name="bottom-orange-wave")
    return slide


def add_footer_badge(slide, kind: str, label: str, x: float):
    # 图标由简单可编辑几何形状重建，不是位图。
    if kind == "leaf":
        add_leaf(slide, x + 0.25, 14.45, 1.05, fill=WHITE)
    elif kind == "cycle":
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.CIRCULAR_ARROW, x + 0.18, 14.35, 0.75, 0.75,
                  fill=WHITE, radius_name="cycle-icon")
    elif kind == "shield":
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.PENTAGON, x + 0.18, 14.32, 0.78, 0.78,
                  fill=GREEN, line=WHITE, width=1.3, radius_name="shield-icon")
        add_text(slide, "✓", x + 0.30, 14.40, 0.54, 0.46, size=13, font="Arial",
                 fill=WHITE, bold=True, align=PP_ALIGN.CENTER)
    else:
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + 0.18, 14.34, 0.76, 0.76,
                  fill=GREEN, line=WHITE, width=1.2, radius_name="globe-icon")
        add_line(slide, x + 0.56, 14.38, x + 0.56, 15.05, fill=WHITE, width=0.8)
        add_line(slide, x + 0.26, 14.71, x + 0.86, 14.71, fill=WHITE, width=0.8)
    add_text(slide, label, x - 0.05, 15.18, 1.55, 0.42, size=7.5, font="Arial",
             fill=WHITE, align=PP_ALIGN.CENTER)


def add_slide_two(prs: Presentation, qr_path: Path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_common_frame(slide)
    add_text(slide, "FUJIAN LANYU NOVA\nMATERIALS CO., LTD.", 5.7, 3.15, 14.2, 1.85,
             size=28, font="Arial", fill=DARK_GREEN, bold=True, anchor=MSO_ANCHOR.TOP)
    add_leaf_divider(slide, 5.75, 5.37, 7.3)
    add_qr(slide, qr_path)

    add_text(slide, "Ethan Shen", 6.15, 6.1, 8.5, 1.45, size=32,
             font="Segoe Script", fill=DARK_GREEN, italic=True)
    add_contact_icon(slide, "☎", 6.4, 8.58)
    add_text(slide, "M.P.   139 4826 7153", 7.6, 8.58, 7.4, 0.9, size=17,
             font="Arial", fill=INK, bold=True)
    add_contact_icon(slide, "✉", 6.4, 9.78)
    add_text(slide, "z.xing@lanyunova.com", 7.6, 9.78, 7.6, 0.9, size=14,
             font="Arial", fill=INK)
    add_contact_icon(slide, "⌖", 6.4, 10.98)
    add_text(slide, "Suite 1208, BlueBay Innovation Center,\nNo. 88 Yunchuang Avenue,\nHaicang District, Xiamen, Fujian, China",
             7.6, 10.85, 8.8, 1.72, size=10.5, font="Arial", fill=INK,
             anchor=MSO_ANCHOR.TOP)
    add_line(slide, 17.05, 6.25, 17.05, 13.75, fill="B7C8BE", width=1.0)

    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 18.2, 6.2, 4.3, 0.72,
              fill=GREEN, radius_name="main-business-pill")
    add_text(slide, "MAIN BUSINESS", 18.2, 6.21, 4.3, 0.68, size=12, font="Arial",
             fill=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_business_list(slide, 18.25, 7.25, english=True)

    add_shape(slide, MSO_AUTO_SHAPE_TYPE.WAVE, 4.28, 13.6, 25.28, 2.95,
              fill=GREEN, radius_name="bottom-green-wave")
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.WAVE, 4.28, 13.55, 17.0, 0.75,
              fill=PALE_ORANGE, radius_name="bottom-orange-wave")
    for kind, label, x in [("leaf", "SUSTAINABLE", 5.6), ("cycle", "INNOVATIVE", 8.0),
                           ("shield", "RELIABLE", 10.4), ("globe", "GLOBAL", 12.8)]:
        add_footer_badge(slide, kind, label, x)
    return slide


def build(source: Path, output: Path) -> None:
    prs = Presentation()
    prs.slide_width = Cm(33.867)
    prs.slide_height = Cm(19.05)
    while prs.slides:
        rel = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(rel)
        del prs.slides._sldIdLst[0]

    with tempfile.TemporaryDirectory(prefix="officeval043-") as temp_dir:
        qr_path = Path(temp_dir) / "qr-source-crop.png"
        with Image.open(source) as image:
            # 只保留参考图中的二维码；其他元素均重建为原生对象。
            image.crop((892, 63, 1074, 245)).save(qr_path)
        add_slide_one(prs, qr_path)
        add_slide_two(prs, qr_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        prs.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    source = args.input_dir / INPUT_NAME
    if not source.is_file() or sha256(source) != INPUT_SHA256:
        raise ValueError("名片参考图缺失或 SHA256 不匹配")
    output = args.output_dir / OUTPUT_NAME
    build(source, output)
    print(output)


if __name__ == "__main__":
    main()
