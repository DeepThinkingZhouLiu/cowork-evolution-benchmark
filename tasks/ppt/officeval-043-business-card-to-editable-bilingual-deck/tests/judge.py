#!/usr/bin/env python3
"""officeval_043 的独立 Harbor Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "lanyu_business_card_editable.pptx"
TASK_ID = "cowork-evo/officeval-043-business-card-to-editable-bilingual-deck"
RUBRIC_VERSION = "omegause-officeval-043-adaptation-v0.1"
JUDGE_VERSION = "officeval-043-adaptation-v0.1"
EMU_PER_CM = 360000.0

SPECS = {
    "R001": ("指定 PPTX 可打开，严格为 2 页 16:9", 4, "hurdle"),
    "R002": ("第 1 页中英文企业与姓名信息完整", 5, "hurdle"),
    "R003": ("第 2 页英文企业与姓名信息完整且无中文", 5, "hurdle"),
    "R004": ("两页联系方式、地址和联系图标完整", 4, "positive"),
    "R005": ("两页主营业务的三组标题与子项完整", 6, "positive"),
    "R006": ("两页外框、分栏和主要文字的几何位置符合参考图", 4, "positive"),
    "R007": ("右上角二维码均为来源图小图片，页面无大图压平", 5, "hurdle"),
    "R008": ("标识、联系图标、业务图标和底部视觉元素均为可编辑对象", 5, "positive"),
    "R009": ("第 2 页底部四个英文品牌标签完整", 3, "positive"),
    "R010": ("LibreOffice 可实际渲染为 2 页，关键文字可提取", 4, "positive"),
    "R011": ("若出现整页背景图、缺少一页或英文页含中文，扣 12 分", -12, "penalty"),
}


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def all_text(slide) -> str:
    return "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text_frame") and shape.has_text_frame)


def has_all(text: str, values: list[str]) -> tuple[bool, list[str]]:
    normalized = normalize(text).lower()
    missing = [value for value in values if normalize(value).lower() not in normalized]
    return not missing, missing


def cm(value: int) -> float:
    return value / EMU_PER_CM


def criterion(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {
        "criterion_id": identifier,
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
        "raw": {"description": description, "weight": weight, "type": kind, "passed": bool(passed)},
    }


def render_check(path: Path) -> tuple[bool, str]:
    with tempfile.TemporaryDirectory(prefix="officeval043-judge-") as temp_dir:
        result = subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", temp_dir, str(path)],
            stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=120,
        )
        pdf = Path(temp_dir) / f"{path.stem}.pdf"
        if result.returncode != 0 or not pdf.is_file():
            return False, f"转 PDF 失败：{result.stdout[-300:]}"
        info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, text=True, timeout=30)
        match = re.search(r"^Pages:\s+(\d+)", info.stdout, re.MULTILINE)
        pages = int(match.group(1)) if match else 0
        text_path = Path(temp_dir) / "rendered.txt"
        subprocess.run(["pdftotext", str(pdf), str(text_path)], check=True, timeout=30)
        rendered = text_path.read_text(errors="replace")
        required = ["沈知行", "Ethan Shen", "MAIN BUSINESS", "139 4826 7153"]
        ok = pages == 2 and all(value.lower() in normalize(rendered).lower() for value in required)
        return ok, f"页数={pages}，关键文字可提取={ok}"


def evaluate(path: Path) -> dict:
    criteria = []
    try:
        prs = Presentation(path)
        loaded = True
    except Exception as exc:  # pragma: no cover - 损坏文件路径
        prs = None
        loaded = False
        load_error = str(exc)

    if not loaded:
        criteria.append(criterion("R001", False, f"打开失败：{load_error}"))
        for identifier in list(SPECS)[1:]:
            criteria.append(criterion(identifier, False, "文件无法解析"))
    else:
        size_ok = abs(cm(prs.slide_width) - 33.867) <= 0.05 and abs(cm(prs.slide_height) - 19.05) <= 0.05
        criteria.append(criterion("R001", len(prs.slides) == 2 and size_ok,
                                   f"页数={len(prs.slides)}，页面={cm(prs.slide_width):.3f}×{cm(prs.slide_height):.3f}cm"))
        slides = list(prs.slides)
        text1 = all_text(slides[0]) if len(slides) > 0 else ""
        text2 = all_text(slides[1]) if len(slides) > 1 else ""

        ok, missing = has_all(text1, ["福建蓝屿新材料科技有限公司", "FUJIAN LANYU NOVA MATERIALS CO., LTD.", "沈知行"])
        criteria.append(criterion("R002", ok, f"第1页缺失={missing}"))
        ok2, missing2 = has_all(text2, ["FUJIAN LANYU NOVA MATERIALS CO., LTD.", "Ethan Shen"])
        chinese2 = re.findall(r"[\u3400-\u9fff]", text2)
        criteria.append(criterion("R003", ok2 and not chinese2,
                                   f"第2页缺失={missing2}，中文字符数={len(chinese2)}"))

        contact1, miss_contact1 = has_all(text1, ["M.P.", "139 4826 7153", "z.xing@lanyunova.com", "蓝湾创新中心12层1208室"])
        contact2, miss_contact2 = has_all(text2, ["M.P.", "139 4826 7153", "z.xing@lanyunova.com", "Suite 1208", "Haicang District"])
        icon_counts = [sum(1 for shape in slide.shapes if shape.name.startswith("contact-")) for slide in slides[:2]]
        criteria.append(criterion("R004", contact1 and contact2 and icon_counts == [3, 3],
                                   f"缺失={miss_contact1 + miss_contact2}，联系图标={icon_counts}"))

        business1, miss_b1 = has_all(text1, ["环保涂层", "水性树脂", "功能助剂", "高阻隔涂层", "绿色材料", "可降解薄膜", "纸塑复合材料", "应用方案", "工业胶黏剂", "包装表面处理"])
        business2, miss_b2 = has_all(text2, ["ECO COATINGS", "Water-based resins", "Functional additives", "Barrier coatings", "GREEN MATERIALS", "Biodegradable films", "Paper-plastic composites", "APPLICATION SOLUTIONS", "Industrial adhesives", "Packaging surface treatment"])
        criteria.append(criterion("R005", business1 and business2, f"业务缺失={miss_b1 + miss_b2}"))

        frames = []
        dividers = []
        for slide in slides[:2]:
            frame = [shape for shape in slide.shapes if shape.name == "business-card-frame"]
            frames.append(bool(frame and 24 <= cm(frame[0].width) <= 26 and 13 <= cm(frame[0].height) <= 15))
            dividers.append(any(abs(cm(shape.height)) >= 6 and cm(shape.width) < 0.2 for shape in slide.shapes))
        criteria.append(criterion("R006", frames == [True, True] and dividers == [True, True],
                                   f"名片外框={frames}，竖向分栏={dividers}"))

        picture_records = []
        flattened = False
        for slide in slides[:2]:
            pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
            good_qr = [shape for shape in pictures if 3 <= cm(shape.width) <= 4 and 3 <= cm(shape.height) <= 4
                       and cm(shape.left) / cm(prs.slide_width) >= 0.7 and cm(shape.top) / cm(prs.slide_height) <= 0.25]
            picture_records.append((len(pictures), len(good_qr)))
            flattened = flattened or any((shape.width * shape.height) / (prs.slide_width * prs.slide_height) >= 0.60 for shape in pictures)
        criteria.append(criterion("R007", picture_records == [(1, 1), (1, 1)] and not flattened,
                                   f"图片/合格二维码={picture_records}，整页大图={flattened}"))

        native_counts = [sum(1 for shape in slide.shapes if shape.shape_type != MSO_SHAPE_TYPE.PICTURE) for slide in slides[:2]]
        named_parts = [sum(1 for shape in slide.shapes if any(token in shape.name for token in ["logo-", "contact-", "business-", "wave", "icon"])) for slide in slides[:2]]
        criteria.append(criterion("R008", all(value >= 50 for value in native_counts) and all(value >= 15 for value in named_parts),
                                   f"原生对象={native_counts}，标识化可编辑部件={named_parts}"))

        labels = ["SUSTAINABLE", "INNOVATIVE", "RELIABLE", "GLOBAL"]
        label_ok, label_missing = has_all(text2, labels)
        criteria.append(criterion("R009", label_ok, f"底部标签缺失={label_missing}"))

        render_ok, render_evidence = render_check(path)
        criteria.append(criterion("R010", render_ok, render_evidence))
        severe = len(prs.slides) != 2 or flattened or bool(chinese2)
        criteria.append(criterion("R011", not severe, f"严重错误={severe}"))

    hurdles = [item["criterion_id"] for item in criteria if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]]
    positive_total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    positive_score = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0 and item["raw"]["passed"])
    penalties = sum(abs(item["raw"]["weight"]) for item in criteria if item["raw"]["weight"] < 0 and not item["raw"]["passed"])
    reward = 0.0 if hurdles else max(0.0, min(1.0, (positive_score - penalties) / positive_total))
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [hashlib.sha256(path.read_bytes()).hexdigest()] if path.is_file() else [],
        "criterion_results": criteria,
        "reward": reward,
        "passed": reward >= 0.8 and not hurdles,
        "hurdle_failures": hurdles,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    output = args.artifact_dir / OUTPUT_NAME
    result = evaluate(output) if output.is_file() else {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": [], "criterion_results": [], "reward": 0.0,
        "passed": False, "hurdle_failures": ["missing_output"],
    }
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.write_text(f"{result['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
