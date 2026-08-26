#!/usr/bin/env python3
"""officeval_046 逐题解法：逐对象重建两张中文科研框架图。"""

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Pt


OUTPUT_NAME = "dual_research_frameworks_editable.pptx"
INPUT_HASHES = {
    "image_001.png": "b09980a366b1de100a43518a4b9cafe7543950af09921b055c64a645d8cd4149",
    "image_002.png": "6b5bf5dacaf7e53c03103df37050e0e25027e21eedc1d06ccdb8965e1b3d7d50",
}
NAVY = "103E77"
BLUE = "2E69B3"
PALE_BLUE = "EAF3FC"
TEAL = "1B8A83"
GREEN = "4C9270"
PALE_GREEN = "EAF5ED"
PURPLE = "7651A6"
PALE_PURPLE = "F1ECF8"
ORANGE = "E79B27"
PALE_ORANGE = "FFF6E7"
RED = "C71919"
INK = "14243A"
GRAY = "EEF1F4"
WHITE = "FFFFFF"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def shape(slide, kind, x: float, y: float, w: float, h: float, *,
          fill: str = WHITE, line: str | None = NAVY, width: float = 1.0,
          dash: bool = False, name: str = ""):
    item = slide.shapes.add_shape(kind, Cm(x), Cm(y), Cm(w), Cm(h))
    item.name = name or f"shape-{len(slide.shapes)}"
    item.fill.solid()
    item.fill.fore_color.rgb = rgb(fill)
    if line is None:
        item.line.fill.background()
    else:
        item.line.color.rgb = rgb(line)
        item.line.width = Pt(width)
        if dash:
            item.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return item


def text(slide, value: str, x: float, y: float, w: float, h: float, *,
         size: float = 12, fill: str = INK, bold: bool = False,
         align=PP_ALIGN.CENTER, font: str = "Microsoft YaHei",
         anchor=MSO_ANCHOR.MIDDLE, vertical: bool = False):
    box = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    box.name = f"text-{value[:20]}"
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Cm(0.03)
    frame.margin_right = Cm(0.03)
    frame.margin_top = Cm(0.02)
    frame.margin_bottom = Cm(0.02)
    frame.vertical_anchor = anchor
    if vertical:
        value = "\n".join(value)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(fill)
    return box


def line(slide, x1: float, y1: float, x2: float, y2: float, *,
         fill: str = NAVY, width: float = 1.0, dash: bool = False,
         arrow: bool = False, both: bool = False, name: str = ""):
    item = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    item.name = name or f"line-{len(slide.shapes)}"
    item.line.color.rgb = rgb(fill)
    item.line.width = Pt(width)
    if dash:
        item.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        item.line.end_arrowhead = True
    if both:
        item.line.begin_arrowhead = True
        item.line.end_arrowhead = True
    return item


def labeled_shape(slide, kind, value: str, x: float, y: float, w: float, h: float, *,
                  fill: str = WHITE, border: str = NAVY, size: float = 12,
                  text_fill: str = INK, bold: bool = False, dash: bool = False,
                  name: str = ""):
    item = shape(slide, kind, x, y, w, h, fill=fill, line=border, dash=dash, name=name)
    frame = item.text_frame
    frame.clear()
    frame.margin_left = Cm(0.08)
    frame.margin_right = Cm(0.08)
    frame.margin_top = Cm(0.03)
    frame.margin_bottom = Cm(0.03)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = value
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(text_fill)
    return item


def icon_drop(slide, x: float, y: float, fill: str = BLUE):
    drop = shape(slide, MSO_AUTO_SHAPE_TYPE.TEAR, x, y, 0.55, 0.82, fill=fill, line=NAVY, name="icon-drop")
    drop.rotation = 0
    return drop


def icon_particles(slide, x: float, y: float, fill: str = BLUE):
    for dx, dy, size in [(0.0, 0.18, .18), (.26, 0.0, .16), (.52, .23, .18), (.24, .45, .14), (.62, .50, .13)]:
        shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + dx, y + dy, size, size, fill=fill, line=None, name="icon-particle")


def icon_shield(slide, x: float, y: float, fill: str = WHITE, border: str = BLUE):
    shape(slide, MSO_AUTO_SHAPE_TYPE.PENTAGON, x, y, 0.7, 0.82, fill=fill, line=border, width=1.2, name="icon-shield")
    text(slide, "✓", x + .12, y + .17, .46, .34, size=13, fill=border, bold=True)


def small_card(slide, value: str, x: float, y: float, icon: str, *, fill: str = WHITE):
    labeled_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, value, x, y, 3.55, 1.35,
                  fill=fill, border="78A5D0", size=10.5, name=f"module-item-{value}")
    if icon == "drop":
        icon_drop(slide, x + .35, y + .23)
    elif icon == "particles":
        icon_particles(slide, x + .27, y + .28)
    elif icon == "layers":
        for index in range(2):
            shape(slide, MSO_AUTO_SHAPE_TYPE.PARALLELOGRAM, x + .25, y + .32 + index * .28, .9, .32,
                  fill=PALE_BLUE if index else BLUE, line=NAVY, width=.6, name="icon-layer")
    elif icon == "coating":
        shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x + .28, y + .45, .9, .34, fill=BLUE, line=NAVY, width=.6, name="icon-coating")
        shape(slide, MSO_AUTO_SHAPE_TYPE.PARALLELOGRAM, x + .38, y + .28, .88, .30, fill=WHITE, line=NAVY, width=.6, name="icon-coating-top")
    elif icon == "shield":
        icon_shield(slide, x + .28, y + .24)
    elif icon == "waves":
        for offset in range(3):
            line(slide, x + .25, y + .45 + offset * .18, x + 1.05, y + .45 + offset * .18, fill=BLUE, width=1)
        for dx, dy, radius in [(.35, .15, .12), (.62, .06, .15), (.91, .2, .1)]:
            shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + dx, y + dy, radius, radius, fill=WHITE, line=BLUE, width=.8, name="icon-bubble")
    elif icon == "sliders":
        for offset in range(3):
            yy = y + .30 + offset * .25
            line(slide, x + .26, yy, x + 1.08, yy, fill=BLUE, width=1)
            shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + .48 + offset * .18, yy - .10, .2, .2, fill=WHITE, line=BLUE, width=.8, name="icon-slider")


def membrane_device(slide):
    # 右侧膜分离设备由原生几何体拼成，不使用设备截图。
    x, y = 23.35, 5.0
    shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + .65, y - .72, .85, .85, fill=GRAY, line="49545E", name="device-gauge")
    line(slide, x + 1.07, y - .30, x + 1.30, y - .62, fill="49545E")
    shape(slide, MSO_AUTO_SHAPE_TYPE.CAN, x, y, 2.2, 1.0, fill="BBC2C7", line="49545E", name="device-top-cap")
    shape(slide, MSO_AUTO_SHAPE_TYPE.CAN, x + .18, y + .55, 1.84, 6.7, fill="D8DDE0", line="49545E", name="device-body")
    for offset in range(1, 9):
        line(slide, x + .25 + offset * .18, y + 1.0, x + .25 + offset * .18, y + 6.8, fill=WHITE, width=.7, name="device-filter-line")
    for yy in [y + 1.35, y + 5.15]:
        shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x - .55, yy, .75, .35, fill="AEB6BC", line="49545E", width=.7, name="device-pipe")
        shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x - .72, yy - .03, .35, .41, fill=GRAY, line="49545E", width=.7, name="device-pipe-mouth")
    shape(slide, MSO_AUTO_SHAPE_TYPE.CAN, x, y + 6.85, 2.2, .8, fill="AEB6BC", line="49545E", name="device-bottom-cap")
    shape(slide, MSO_AUTO_SHAPE_TYPE.CAN, x + .65, y + 7.45, .9, .5, fill="9DA6AD", line="49545E", name="device-bottom-neck")
    for dx in [.18, .92, 1.66]:
        shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x + dx, y + 7.85, .22, .7, fill="8C969E", line="49545E", name="device-leg")


def slide_membrane(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(WHITE)
    # 箱体正面、左侧指标面和右侧设备面。
    shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 4.45, 2.6, 18.25, 14.65, fill="FDFEFE", line=NAVY, width=1.0, name="front-panel")
    shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 2.0, 2.6, 2.45, 14.65, fill="F5F9FC", line=NAVY, width=1.0, name="metric-panel")
    shape(slide, MSO_AUTO_SHAPE_TYPE.PARALLELOGRAM, 22.7, 2.6, 3.45, 14.65, fill="F1F5F8", line=NAVY, width=1.0, name="device-panel")
    tabs = [(4.4, "运行系统", "DCECF8"), (9.0, "膜材料", "E9ECEF"), (13.6, "污染层生长", "EEF4EC"), (18.4, "处理系统", "DCECF8")]
    for x, label, fill in tabs:
        labeled_shape(slide, MSO_AUTO_SHAPE_TYPE.PARALLELOGRAM, label, x, .95, 4.7, 1.65,
                      fill=fill, border=NAVY, size=16, text_fill=NAVY, bold=True, name=f"top-tab-{label}")

    # 左侧三段指标栏。
    for yy in [7.2, 12.0]:
        line(slide, 2.0, yy, 4.45, yy, fill="829CB5", dash=True)
    shape(slide, MSO_AUTO_SHAPE_TYPE.ARC, 2.60, 3.35, 1.05, .75, fill=WHITE, line=NAVY, width=1.2, name="pressure-gauge")
    line(slide, 3.12, 3.74, 3.55, 3.36, fill=NAVY, width=1.2)
    text(slide, "跨膜压差", 2.12, 4.18, 2.18, .55, size=11.5, bold=True)
    line(slide, 3.22, 4.9, 3.22, 6.45, fill=BLUE, width=1.5, arrow=True)
    icon_drop(slide, 2.66, 8.05)
    icon_particles(slide, 2.83, 8.55)
    text(slide, "污染迁移速率", 2.05, 9.18, 2.3, .78, size=10.5, bold=True)
    line(slide, 3.22, 10.1, 3.22, 11.55, fill=BLUE, width=1.5, arrow=True)
    shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 2.57, 13.0, 1.2, 1.15, fill=PALE_BLUE, line=NAVY, name="resistance-particle-box")
    for row in range(4):
        for col in range(5):
            shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 2.66 + col * .20, 13.12 + row * .22, .09, .09, fill=BLUE, line=None, name="resistance-dot")
    text(slide, "传质阻力", 2.1, 14.45, 2.2, .55, size=11.5, bold=True)

    nodes = [(5.45, "膜组件"), (9.65, "污染沉积层"), (13.9, "表面相互作用"), (18.25, "溶液性质")]
    for index, (x, label) in enumerate(nodes):
        labeled_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, label, x, 3.45, 3.1, 1.1,
                      fill=WHITE, border=BLUE, size=11.5, bold=True, name=f"upper-node-{index+1}")
        if index:
            line(slide, x - .85, 4.0, x - .12, 4.0, fill=NAVY, width=1.3, arrow=True)
        line(slide, x + 1.55, 4.55, x + 1.55, 5.15, fill=BLUE, dash=True)
    line(slide, 7.0, 5.15, 19.8, 5.15, fill=BLUE, dash=True)
    line(slide, 13.4, 5.15, 13.4, 5.55, fill=BLUE, dash=True, arrow=True)
    labeled_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                  "膜污染演化是影响通量衰减的重要因素之一", 8.45, 5.55, 10.8, 1.05,
                  fill=PALE_BLUE, border=BLUE, size=13.5, text_fill=NAVY, bold=True, name="evolution-note")
    line(slide, 13.85, 6.6, 13.85, 7.3, fill=NAVY, width=1.3, arrow=True)
    labeled_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                  "精细调控膜表面结构与运行参数\n优化污染控制状态", 7.55, 7.35, 12.8, 2.0,
                  fill=PALE_ORANGE, border=ORANGE, size=15, text_fill=INK, bold=True, dash=True, name="control-target")
    # 两个下方模块和 8 个内嵌小卡。
    labeled_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, "", 5.25, 10.15, 7.65, 4.9,
                  fill=PALE_BLUE, border=BLUE, name="interface-module")
    text(slide, "界面相容优化", 5.6, 10.2, 6.95, .65, size=15, fill=NAVY, bold=True)
    labeled_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, "", 14.1, 10.15, 7.65, 4.9,
                  fill=WHITE, border=BLUE, name="operation-module")
    text(slide, "运行抑污控制", 14.45, 10.2, 6.95, .65, size=15, fill=NAVY, bold=True)
    left_items = [(5.45, 11.0, "增强表面\n亲水性", "drop"), (9.2, 11.0, "改进\n孔结构", "particles"),
                  (5.45, 12.55, "膜基材", "layers"), (9.2, 12.55, "表面涂层", "coating")]
    right_items = [(14.3, 11.0, "抑制浓差\n极化", "particles"), (18.05, 11.0, "降低不可\n逆污染", "shield"),
                   (14.3, 12.55, "气水反冲", "waves"), (18.05, 12.55, "运行参数", "sliders")]
    for x, y, label, icon in left_items + right_items:
        small_card(slide, label, x, y, icon, fill=WHITE)
    line(slide, 13.85, 9.35, 13.85, 9.75, fill=NAVY)
    line(slide, 9.05, 9.75, 18.0, 9.75, fill=NAVY)
    line(slide, 9.05, 9.75, 9.05, 10.15, fill=NAVY, arrow=True)
    line(slide, 18.0, 9.75, 18.0, 10.15, fill=NAVY, arrow=True)
    line(slide, 9.05, 15.05, 9.05, 15.35, fill=NAVY)
    line(slide, 18.0, 15.05, 18.0, 15.35, fill=NAVY)
    line(slide, 9.05, 15.35, 18.0, 15.35, fill=NAVY)
    line(slide, 13.55, 15.35, 13.55, 15.7, fill=NAVY, arrow=True)
    labeled_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                  "通量衰减减缓   ·   清洗周期延长   ·   分离效率提升", 6.75, 15.75, 13.6, 1.05,
                  fill=PALE_BLUE, border=BLUE, size=13.5, text_fill=NAVY, bold=True, name="membrane-result")
    membrane_device(slide)
    text(slide, "降低膜污染并提升分离效率", 27.1, 4.4, 1.0, 10.5, size=18.5, fill=RED, bold=True, vertical=True)
    text(slide, "图2  膜分离水处理系统污染控制与通量提升研究", 7.5, 18.55, 15.0, .8, size=15.5, bold=True)


def header_stage(slide, x: float, value: str, fill: str, icon: str):
    labeled_shape(slide, MSO_AUTO_SHAPE_TYPE.CHEVRON, value, x, .92, 6.1, 1.55,
                  fill=fill, border=fill, size=14.5, text_fill=WHITE, bold=True, name=f"pipeline-{value}")
    shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + .28, 1.10, .95, .95, fill=WHITE, line=None, name=f"pipeline-icon-bg-{icon}")
    text(slide, icon, x + .38, 1.20, .75, .70, size=18, fill=fill, bold=True)


def stacked_module(slide, x: float, y: float, title_value: str, items: list[str], *, theme: str, pale: str):
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, 5.1, 3.65, fill=WHITE, line=theme, name=f"module-{title_value}")
    labeled_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, title_value, x + .25, y + .25, 4.6, .82,
                  fill=theme, border=theme, size=12.5, text_fill=WHITE, bold=True, name=f"module-title-{title_value}")
    for index, item in enumerate(items):
        labeled_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, item, x + .52, y + 1.38 + index * 1.02,
                      4.05, .78, fill=pale, border=theme, size=11.5, name=f"module-item-{item}")


def chart_bars(slide, x: float, y: float):
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, 5.35, 4.05, fill=WHITE, line="AAB3BD", name="feature-importance-chart")
    text(slide, "特征重要性（示例）", x + .5, y + .18, 4.35, .38, size=9.5)
    line(slide, x + 1.35, y + .8, x + 1.35, y + 3.45, fill="555555", width=.7)
    line(slide, x + 1.35, y + 3.45, x + 4.95, y + 3.45, fill="555555", width=.7)
    labels = ["心率变异性", "睡眠时长", "日均步数", "体重指数", "血压波动", "吸烟状态"]
    widths = [3.05, 2.65, 2.25, 1.75, 1.35, .85]
    for index, (label, width) in enumerate(zip(labels, widths)):
        yy = y + .85 + index * .38
        text(slide, label, x + .10, yy, 1.15, .30, size=6.8, align=PP_ALIGN.RIGHT)
        shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x + 1.4, yy + .03, width, .22, fill=BLUE, line=None, name="chart-bar")
    text(slide, "平均绝对SHAP值", x + 1.65, y + 3.48, 2.9, .35, size=7.5)


def chart_lines(slide, x: float, y: float):
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, 5.7, 4.05, fill=WHITE, line="AAB3BD", name="risk-trend-chart")
    text(slide, "风险随时间变化（示例）", x + .55, y + .18, 4.7, .38, size=9.5)
    line(slide, x + 1.1, y + .8, x + 1.1, y + 3.45, fill="555555", width=.7)
    line(slide, x + 1.1, y + 3.45, x + 5.05, y + 3.45, fill="555555", width=.7)
    series = [(GREEN, [2.45, 2.1, 2.25, 1.8, 2.0, 1.55, 1.75]),
              (ORANGE, [2.95, 2.78, 2.88, 2.5, 2.65, 2.2, 2.35]),
              ("D73A32", [3.28, 3.18, 3.12, 2.98, 2.92, 2.75, 2.62])]
    for tone, values in series:
        for idx in range(len(values) - 1):
            line(slide, x + 1.2 + idx * .55, y + values[idx], x + 1.75 + idx * .55, y + values[idx + 1], fill=tone, width=1.0, name="trend-segment")
    for index, (tone, label) in enumerate([(GREEN, "低风险"), (ORANGE, "中风险"), ("D73A32", "高风险")]):
        line(slide, x + 4.38, y + .85 + index * .32, x + 4.72, y + .85 + index * .32, fill=tone, width=1.2)
        text(slide, label, x + 4.75, y + .73 + index * .32, .8, .28, size=6.8, align=PP_ALIGN.LEFT)
    text(slide, "随访时间（月）", x + 2.25, y + 3.48, 2.4, .35, size=7.5)


def bottom_pill(slide, x: float, label: str, fill: str, icon: str):
    labeled_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, label, x, 16.55, 5.45, 1.75,
                  fill=fill, border=fill, size=14.5, text_fill=WHITE, bold=True, name=f"bottom-pill-{label}")
    text(slide, icon, x + .38, 16.86, .72, .72, size=18, fill=WHITE, bold=True)


def slide_chronic(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(WHITE)
    for x, label, tone, icon in [(1.25, "数据获取与治理", TEAL, "≈"), (8.25, "特征提取与建模", BLUE, "∷"),
                                  (15.25, "模型评估与解释", PURPLE, "⌕"), (22.25, "干预建议生成", ORANGE, "+")]:
        header_stage(slide, x, label, tone, icon)

    stacked_module(slide, 1.35, 3.2, "传统机器学习", ["随机森林", "支持向量机"], theme=TEAL, pale=PALE_GREEN)
    stacked_module(slide, 1.35, 7.35, "时序深度学习", ["生理信号", "生活行为序列"], theme=BLUE, pale=PALE_BLUE)
    labeled_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                  "利用时间依赖建模捕捉\n风险状态的动态变化", 1.15, 11.55, 5.8, 1.95,
                  fill=PALE_ORANGE, border=ORANGE, size=11.5, bold=True, dash=True, name="temporal-note")
    line(slide, 4.05, 13.55, 4.05, 14.25, fill=GREEN, width=1.8, arrow=True)
    labeled_shape(slide, MSO_AUTO_SHAPE_TYPE.TRAPEZOID, "健康数据", 2.45, 14.25, 3.2, 1.0,
                  fill=PALE_GREEN, border=GREEN, size=13, bold=True, name="health-data")

    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 7.95, 3.15, 13.8, 13.0,
          fill=WHITE, line=BLUE, dash=True, name="central-analysis-frame")
    labeled_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, "风险等级", 13.55, 5.85, 3.25, 3.25,
                  fill=BLUE, border=NAVY, size=18, text_fill=WHITE, bold=True, name="risk-level")
    satellites = [
        (10.15, 3.45, 2.65, 2.35, "连续指标", PALE_GREEN, GREEN, MSO_AUTO_SHAPE_TYPE.OVAL),
        (17.25, 3.5, 2.55, 2.3, "症状量表", PALE_PURPLE, PURPLE, MSO_AUTO_SHAPE_TYPE.HEXAGON),
        (17.25, 8.25, 2.55, 2.25, "异常事件识别", "E3F4F6", TEAL, MSO_AUTO_SHAPE_TYPE.DIAMOND),
        (10.15, 8.25, 2.65, 2.25, "生活方式因素", "FFF3D7", "C79523", MSO_AUTO_SHAPE_TYPE.HEXAGON),
    ]
    for x, y, w, h, label, pale, tone, kind in satellites:
        labeled_shape(slide, kind, label, x, y, w, h, fill=pale, border=tone, size=10.5, bold=True, name=f"risk-factor-{label}")
        line(slide, 15.17, 7.47, x + w / 2, y + h / 2, fill=BLUE, width=1.0, both=True, name="risk-bidirectional")
    tags = [(13.95, 4.0, "AUC"), (19.65, 6.7, "SHAP"), (13.55, 9.75, "随访评分"), (8.65, 6.7, "F1")]
    for x, y, label in tags:
        labeled_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, label, x, y, 1.55, .72,
                      fill=PALE_BLUE, border=BLUE, size=10.5, name=f"risk-tag-{label}")
        line(slide, 15.17, 7.47, x + .77, y + .36, fill=BLUE, width=1.0, both=True, name="risk-bidirectional")
    chart_bars(slide, 8.45, 11.05)
    chart_lines(slide, 14.2, 11.05)
    text(slide, "临床表型解析与功能关联", 10.8, 15.5, 8.0, .55, size=14.5, fill=BLUE, bold=True)

    stacked_module(slide, 23.15, 3.2, "训练策略", ["交叉验证", "外部测试"], theme=PURPLE, pale=PALE_PURPLE)
    stacked_module(slide, 23.15, 7.35, "模型泛化能力", ["稳定可靠", "参考基准"], theme=BLUE, pale=PALE_BLUE)
    labeled_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                  "采用ROC曲线与校准曲线\n综合评估模型判别与泛化性能", 22.85, 11.55, 5.9, 1.95,
                  fill=PALE_ORANGE, border=ORANGE, size=10.5, dash=True, name="roc-note")
    line(slide, 25.8, 13.55, 25.8, 14.25, fill=PURPLE, width=1.8, arrow=True)
    labeled_shape(slide, MSO_AUTO_SHAPE_TYPE.TRAPEZOID, "随访评估", 24.2, 14.25, 3.2, 1.0,
                  fill=PALE_PURPLE, border=PURPLE, size=13, bold=True, name="followup-evaluation")

    bottom_pill(slide, 1.15, "构建预测模型", GREEN, "◎")
    bottom_pill(slide, 12.25, "模型验证与整合", BLUE, "✓")
    bottom_pill(slide, 23.4, "可视化趋势曲线", PURPLE, "╱")
    line(slide, 12.2, 17.42, 7.0, 17.42, fill=GREEN, width=2.0, arrow=True)
    line(slide, 17.7, 17.42, 23.35, 17.42, fill=PURPLE, width=2.0, arrow=True)
    text(slide, "图1  多源生理信号慢病风险预测与个体化干预研究框架", 6.5, 18.72, 17.0, .7, size=14, bold=True)


def build(output: Path) -> None:
    prs = Presentation()
    prs.slide_width = Cm(30.0)
    prs.slide_height = Cm(20.0)
    while prs.slides:
        rid = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(rid)
        del prs.slides._sldIdLst[0]
    slide_membrane(prs)
    slide_chronic(prs)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    for name, expected in INPUT_HASHES.items():
        source = args.input_dir / name
        if not source.is_file() or sha256(source) != expected:
            raise ValueError(f"参考图 {name} 缺失或 SHA256 不匹配")
    output = args.output_dir / OUTPUT_NAME
    build(output)
    print(output)


if __name__ == "__main__":
    main()
