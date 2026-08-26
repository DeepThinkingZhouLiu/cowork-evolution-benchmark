#!/usr/bin/env python3
"""officeval_046 两页可编辑科研框架图的独立 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "dual_research_frameworks_editable.pptx"
TASK_ID = "cowork-evo/officeval-046-image-to-editable-dual-research-diagrams"
RUBRIC_VERSION = "omegause-officeval-046-adaptation-v0.1"
JUDGE_VERSION = "officeval-046-adaptation-v0.1"
EMU_PER_CM = 360000.0

SPECS = {
    "R001": ("指定 PPTX 可打开，严格为 2 页 30×20 cm", 4, "hurdle"),
    "R002": ("第 1 页膜污染演化、调控、两个子模块与结果文字完整", 7, "hurdle"),
    "R003": ("第 1 页顶部分栏、左侧指标栏、主流程和两个模块结构完整", 6, "positive"),
    "R004": ("第 1 页膜分离设备为多个原生形状组成", 4, "positive"),
    "R005": ("第 2 页四阶段、左右模块、中心风险结构与底部结果文字完整", 7, "hurdle"),
    "R006": ("第 2 页风险等级、8 个周边标签和连接关系完整", 6, "positive"),
    "R007": ("第 2 页条形图和三糴风险折线图为原生可编辑对象", 5, "positive"),
    "R008": ("两页均由大量原生文字、形状与连线组成，不含位图", 6, "hurdle"),
    "R009": ("两页主体占满页面且无对象越界", 3, "positive"),
    "R010": ("LibreOffice 可实际渲染为 2 页，两个图题与关键节点可读", 4, "positive"),
    "R011": ("若缺少一页、出现大图压平或缺少一个核心框架，扣 15 分", -15, "penalty"),
}

PAGE1_TEXT = [
    "运行系统", "膜材料", "污染层生长", "处理系统", "跨膜压差", "污染迁移速率", "传质阻力",
    "膜组件", "污染沉积层", "表面相互作用", "溶液性质", "膜污染演化是影响通量衰减的重要因素之一",
    "精细调控膜表面结构与运行参数", "优化污染控制状态", "界面相容优化", "运行抑污控制",
    "增强表面", "改进", "膜基材", "表面涂层", "抑制浓差", "降低不可", "气水反冲", "运行参数",
    "通量衰减减缓", "清洗周期延长", "分离效率提升", "膜分离水处理系统污染控制与通量提升研究",
]
PAGE2_TEXT = [
    "数据获取与治理", "特征提取与建模", "模型评估与解释", "干预建议生成", "传统机器学习", "随机森林", "支持向量机",
    "时序深度学习", "生理信号", "生活行为序列", "风险等级", "连续指标", "症状量表", "异常事件识别", "生活方式因素",
    "AUC", "SHAP", "随访评分", "F1", "特征重要性", "风险随时间变化", "低风险", "中风险", "高风险", "临床表型解析与功能关联",
    "训练策略", "交叉验证", "外部测试", "模型泛化能力", "稳定可靠", "参考基准", "ROC曲线", "健康数据", "随访评估",
    "构建预测模型", "模型验证与整合", "可视化趋势曲线", "多源生理信号慢病风险预测与个体化干预研究框架",
]


def cm(value: int) -> float:
    return value / EMU_PER_CM


def normalize(value: str) -> str:
    return re.sub(r"\s+", "", value).lower()


def slide_text(slide) -> str:
    return "\n".join(shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False))


def missing(text: str, expected: list[str]) -> list[str]:
    content = normalize(text)
    return [value for value in expected if normalize(value) not in content]


def criterion(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {"criterion_id": identifier, "score": 1.0 if passed else 0.0,
            "evidence": evidence, "method": "deterministic",
            "raw": {"description": description, "weight": weight, "type": kind, "passed": bool(passed)}}


def render(path: Path) -> tuple[bool, str]:
    with tempfile.TemporaryDirectory(prefix="officeval046-judge-") as temp_dir:
        completed = subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", temp_dir, str(path)],
                                   stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=150)
        pdf = Path(temp_dir) / f"{path.stem}.pdf"
        if completed.returncode != 0 or not pdf.is_file():
            return False, f"转 PDF 失败：{completed.stdout[-300:]}"
        info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, text=True, timeout=30).stdout
        page_match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
        pages = int(page_match.group(1)) if page_match else 0
        txt = Path(temp_dir) / "out.txt"
        subprocess.run(["pdftotext", str(pdf), str(txt)], check=True, timeout=30)
        rendered = normalize(txt.read_text(errors="replace"))
        keys = ["膜污染演化", "风险等级", "图1", "图2"]
        ok = pages == 2 and all(normalize(key) in rendered for key in keys)
        return ok, f"页数={pages}，关键文本可读={ok}"


def evaluate(path: Path) -> dict:
    criteria = []
    try:
        prs = Presentation(path)
    except Exception as exc:
        for identifier in SPECS:
            criteria.append(criterion(identifier, False, f"文件无法打开：{exc}"))
        return result(path, criteria)

    size_ok = abs(cm(prs.slide_width) - 30.0) <= .03 and abs(cm(prs.slide_height) - 20.0) <= .03
    criteria.append(criterion("R001", len(prs.slides) == 2 and size_ok,
                               f"页数={len(prs.slides)}，页面={cm(prs.slide_width):.2f}×{cm(prs.slide_height):.2f}cm"))
    slides = list(prs.slides)
    page1 = slides[0] if slides else None
    page2 = slides[1] if len(slides) > 1 else None
    text1 = slide_text(page1) if page1 else ""
    text2 = slide_text(page2) if page2 else ""
    miss1 = missing(text1, PAGE1_TEXT)
    miss2 = missing(text2, PAGE2_TEXT)
    criteria.append(criterion("R002", not miss1, f"第1页关键文本缺失={miss1}"))

    if page1:
        counts1 = {prefix: sum(1 for shape in page1.shapes if shape.name.startswith(prefix))
                   for prefix in ["top-tab-", "upper-node-", "module-item-", "interface-module", "operation-module"]}
        panel_count = sum(1 for shape in page1.shapes if shape.name in {"front-panel", "metric-panel", "device-panel"})
        structure1 = counts1["top-tab-"] == 4 and counts1["upper-node-"] == 4 and counts1["module-item-"] == 8 and panel_count == 3
    else:
        counts1, panel_count, structure1 = {}, 0, False
    criteria.append(criterion("R003", structure1, f"结构计数={counts1}，箱体面={panel_count}"))
    device_parts = sum(1 for shape in page1.shapes if shape.name.startswith("device-") or shape.name == "pressure-gauge") if page1 else 0
    criteria.append(criterion("R004", device_parts >= 18, f"设备原生部件={device_parts}"))

    criteria.append(criterion("R005", not miss2, f"第2页关键文本缺失={miss2}"))
    if page2:
        pipelines = sum(1 for shape in page2.shapes if shape.name.startswith("pipeline-") and not shape.name.startswith("pipeline-icon"))
        factors = sum(1 for shape in page2.shapes if shape.name.startswith("risk-factor-"))
        tags = sum(1 for shape in page2.shapes if shape.name.startswith("risk-tag-"))
        links = sum(1 for shape in page2.shapes if shape.name == "risk-bidirectional")
        center = sum(1 for shape in page2.shapes if shape.name == "risk-level")
        risk_ok = pipelines == 4 and factors == 4 and tags == 4 and links == 8 and center == 1
    else:
        pipelines = factors = tags = links = center = 0; risk_ok = False
    criteria.append(criterion("R006", risk_ok,
                               f"阶段={pipelines}，因素={factors}，指标标签={tags}，连接={links}，中心={center}"))
    bars = sum(1 for shape in page2.shapes if shape.name == "chart-bar") if page2 else 0
    trends = sum(1 for shape in page2.shapes if shape.name == "trend-segment") if page2 else 0
    charts = sum(1 for shape in page2.shapes if shape.name in {"feature-importance-chart", "risk-trend-chart"}) if page2 else 0
    criteria.append(criterion("R007", bars >= 6 and trends >= 18 and charts == 2,
                               f"图表框={charts}，条形={bars}，折线段={trends}"))

    picture_counts = [sum(1 for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE) for slide in slides]
    native_counts = [len(slide.shapes) - picture_counts[index] for index, slide in enumerate(slides)]
    editable_ok = len(picture_counts) == 2 and picture_counts == [0, 0] and native_counts[0] >= 120 and native_counts[1] >= 100
    criteria.append(criterion("R008", editable_ok, f"图片={picture_counts}，原生对象={native_counts}"))

    inside = []
    for slide in slides:
        inside.append(all(shape.left >= -1000 and shape.top >= -1000 and shape.left + shape.width <= prs.slide_width + 1000 and shape.top + shape.height <= prs.slide_height + 1000 for shape in slide.shapes))
    criteria.append(criterion("R009", inside == [True, True], f"对象全部页内={inside}"))
    render_ok, render_evidence = render(path)
    criteria.append(criterion("R010", render_ok, render_evidence))
    severe = len(prs.slides) != 2 or any(picture_counts) or bool(miss1) or bool(miss2)
    criteria.append(criterion("R011", not severe, f"严重错误={severe}"))
    return result(path, criteria)


def result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [item["criterion_id"] for item in criteria if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0 and item["raw"]["passed"])
    penalties = sum(abs(item["raw"]["weight"]) for item in criteria if item["raw"]["weight"] < 0 and not item["raw"]["passed"])
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalties) / total))
    return {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
            "artifact_sha256": [hashlib.sha256(path.read_bytes()).hexdigest()] if path.is_file() else [],
            "criterion_results": criteria, "reward": reward,
            "passed": reward >= .8 and not hurdles, "hurdle_failures": hurdles}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    output = args.artifact_dir / OUTPUT_NAME
    evaluation = evaluate(output) if output.is_file() else {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": [], "criterion_results": [], "reward": 0.0,
        "passed": False, "hurdle_failures": ["missing_output"]}
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
