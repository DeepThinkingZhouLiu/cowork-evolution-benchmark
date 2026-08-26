#!/usr/bin/env python3
"""officeval_051 逐题解法：将异常检测框架图重建为全可编辑 PPT。"""

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Pt


INPUT_NAME = "Image 2.png"
INPUT_SHA256 = "ef53bc48e47d05feab61746067ff5d98582653a4c9ec53ce111def01e5200eba"
OUTPUT_NAME = "anomaly_detection_editable_flow.pptx"
TEAL = "159B98"
PALE_TEAL = "EFF9F7"
PURPLE = "8B65C2"
PALE_PURPLE = "F7F2FC"
ORANGE = "F07817"
PALE_ORANGE = "FFF7ED"
GREEN = "6DB33F"
PALE_GREEN = "F4FAEE"
PINK = "E45A88"
PALE_PINK = "FFF5F8"
BLUE = "2E79C7"
PALE_BLUE = "F1F7FD"
INK = "181818"
GRAY = "9A9A9A"
WHITE = "FFFFFF"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def shape(slide, kind, x: float, y: float, w: float, h: float, *,
          fill: str = WHITE, border: str | None = INK, width: float = 1.0,
          dash: bool = False, name: str = ""):
    item = slide.shapes.add_shape(kind, Cm(x), Cm(y), Cm(w), Cm(h))
    item.name = name or f"shape-{len(slide.shapes)}"
    item.fill.solid(); item.fill.fore_color.rgb = rgb(fill)
    if border is None:
        item.line.fill.background()
    else:
        item.line.color.rgb = rgb(border); item.line.width = Pt(width)
        if dash:
            item.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return item


def text(slide, value: str, x: float, y: float, w: float, h: float, *,
         size: float = 12, fill: str = INK, bold: bool = False,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
         font: str = "Arial", name: str = ""):
    box = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    box.name = name or f"text-{value[:24]}"
    frame = box.text_frame; frame.clear()
    frame.margin_left = Cm(.03); frame.margin_right = Cm(.03)
    frame.margin_top = Cm(.02); frame.margin_bottom = Cm(.02)
    frame.vertical_anchor = anchor
    paragraph = frame.paragraphs[0]; paragraph.alignment = align; paragraph.space_after = Pt(0)
    run = paragraph.add_run(); run.text = value
    run.font.name = font; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = rgb(fill)
    return box


def labeled(slide, value: str, x: float, y: float, w: float, h: float, *,
            fill: str = WHITE, border: str = INK, size: float = 11,
            bold: bool = False, dash: bool = False, name: str = ""):
    item = shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h,
                 fill=fill, border=border, dash=dash, name=name)
    frame = item.text_frame; frame.clear(); frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = Cm(.06); frame.margin_right = Cm(.06)
    paragraph = frame.paragraphs[0]; paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run(); run.text = value
    run.font.name = "Arial"; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = rgb(INK)
    return item


def line(slide, x1: float, y1: float, x2: float, y2: float, *,
         fill: str = INK, width: float = 1.0, dash: bool = False,
         arrow: bool = False, name: str = ""):
    item = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    item.name = name or f"line-{len(slide.shapes)}"
    item.line.color.rgb = rgb(fill); item.line.width = Pt(width)
    if dash:
        item.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        item.line.end_arrowhead = True
    return item


def elbow(slide, points: list[tuple[float, float]], *, fill: str = INK,
          width: float = 1.0, dash: bool = False, name: str = ""):
    for index, ((x1, y1), (x2, y2)) in enumerate(zip(points, points[1:])):
        line(slide, x1, y1, x2, y2, fill=fill, width=width, dash=dash,
             arrow=index == len(points) - 2, name=f"{name}-segment-{index+1}")


def add_multivariate_card(slide):
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, .6, 5.0, 5.0, 6.2,
          fill=PALE_TEAL, border=TEAL, width=1.2, name="multivariate-card")
    text(slide, "Multivariate\nTime Series", 1.25, 5.45, 3.7, 1.2,
         size=14.5, fill=TEAL, bold=True, name="multivariate-title")
    # 灰色虚线网格与四条独立可编辑折线。
    for index in range(6):
        x = 1.35 + index * .62
        line(slide, x, 7.0, x, 10.45, fill="BFC4C7", width=.7, dash=True, name="timeseries-grid")
    series = [
        (BLUE, [8.0, 7.5, 7.7, 8.2, 8.05, 7.65, 7.72]),
        (ORANGE, [8.15, 8.55, 8.6, 8.2, 8.55, 8.45, 7.55]),
        (PURPLE, [9.0, 8.65, 8.75, 9.15, 8.8, 8.7, 9.05]),
        ("3F9B4B", [10.0, 9.7, 9.95, 9.65, 9.85, 9.45, 9.9]),
    ]
    for tone, values in series:
        for index in range(len(values) - 1):
            line(slide, 1.35 + index * .62, values[index], 1.97 + index * .62, values[index + 1],
                 fill=tone, width=1.4, name="timeseries-line")
    line(slide, 1.25, 10.5, 4.85, 10.5, fill=INK, width=.8)


def add_sequence_encoder(slide):
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 8.35, 1.55, 6.35, 4.9,
          fill=WHITE, border=TEAL, width=1.1, dash=True, name="sequence-encoder-frame")
    text(slide, "1. Sequence Encoder", 8.45, .65, 6.0, .7, size=15.5, fill=TEAL, bold=True)
    labeled(slide, "Gated Recurrent\nUnit", 8.9, 2.05, 5.2, 1.7,
            fill=PALE_TEAL, border=TEAL, size=13, bold=True, name="gru-1")
    labeled(slide, "Gated Recurrent\nUnit", 8.9, 4.15, 5.2, 1.7,
            fill=PALE_TEAL, border=TEAL, size=13, bold=True, name="gru-2")


def add_pattern_encoder(slide):
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 8.35, 8.05, 6.0, 5.25,
          fill=PALE_PURPLE, border=PURPLE, width=1.1, dash=True, name="pattern-encoder-frame")
    text(slide, "2. Pattern Encoder", 8.55, 8.2, 5.6, .65, size=14.5, fill=PURPLE, bold=True)
    tones = ["E8DDF7", "B99ADD", "F9F6FC"]
    for row in range(3):
        for col in range(5):
            shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 9.75 + col * .75, 9.45 + row * .72, .75, .72,
                  fill=tones[(row + col * 2) % 3], border=PURPLE, width=.6, name="embedding-cell")
    text(slide, "Temporal\nEmbedding", 10.0, 11.75, 3.8, 1.1, size=11.5, bold=True)


def add_concat_icon(slide, x: float, y: float):
    for row in range(4):
        for col in range(2):
            shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x + col * .55, y + row * .58, .5, .53,
                  fill=PALE_ORANGE, border=ORANGE, width=.6, name="concat-cell")
        line(slide, x + .50, y + .26 + row * .58, x + .57, y + .26 + row * .58, fill=ORANGE, width=.6, name="concat-join")


def add_attention_icon(slide, x: float, y: float):
    nodes = [(x + .95, y + .12), (x + .25, y + 1.05), (x + .95, y + 1.05), (x + 1.65, y + 1.05)]
    for nx, ny in nodes:
        shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, nx, ny, .3, .3, fill=PALE_ORANGE, border=ORANGE, width=.8, name="attention-node")
    for target in nodes[1:]:
        line(slide, nodes[0][0] + .15, nodes[0][1] + .30, target[0] + .15, target[1], fill=ORANGE, width=.8, dash=True, name="attention-link")
        line(slide, target[0] + .15, target[1] + .30, target[0] + .15, y + 1.65, fill=ORANGE, width=.8, name="attention-stem")
    shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x + .05, y + 1.65, 2.05, .35, fill="F8D8B5", border=ORANGE, width=.7, name="attention-output")


def add_network_icon(slide, x: float, y: float):
    left = [(x, y + .2), (x, y + .95), (x, y + 1.7)]
    middle = [(x + 1.05, y + .55)]
    right = [(x + 2.0, y + .35), (x + 2.0, y + 1.45)]
    for start in left:
        for end in middle:
            line(slide, start[0] + .28, start[1] + .14, end[0] + .14, end[1] + .14, fill=ORANGE, width=.7, name="network-link")
    for start in middle:
        for end in right:
            line(slide, start[0] + .28, start[1] + .14, end[0] + .14, end[1] + .14, fill=ORANGE, width=.7, name="network-link")
    for nx, ny in left + middle + right:
        shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, nx, ny, .3, .3, fill=PALE_ORANGE, border=ORANGE, width=.7, name="network-node")


def add_norm_icon(slide, x: float, y: float):
    line(slide, x, y + 1.75, x, y, fill=INK, width=.9)
    line(slide, x, y + 1.75, x + 2.0, y + 1.75, fill=INK, width=.9, arrow=True)
    points = [(0, 1.45), (.35, 1.35), (.65, .85), (.95, .35), (1.25, .45), (1.55, 1.15), (1.9, 1.35)]
    for first, second in zip(points, points[1:]):
        line(slide, x + first[0], y + first[1], x + second[0], y + second[1], fill=ORANGE, width=1.1, name="normalization-curve")


def add_fusion(slide):
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 16.15, 1.65, 15.05, 7.8,
          fill=WHITE, border=ORANGE, width=1.1, dash=True, name="fusion-frame")
    text(slide, "3. Fusion & Representation Learning", 17.85, .68, 11.7, .75,
         size=15.5, fill=ORANGE, bold=True)
    cards = [(16.65, "Concatenate"), (20.35, "Self-Attention\nBlock"),
             (24.1, "Feed-Forward\nNetwork"), (27.85, "Layer\nNormalization")]
    for index, (x, label) in enumerate(cards):
        labeled(slide, "", x, 2.55, 3.0, 5.7, fill=PALE_ORANGE, border=ORANGE,
                size=11, name=f"fusion-card-{index+1}")
        text(slide, label, x + .18, 2.85, 2.64, 1.0, size=10.5, bold=False,
             name=f"fusion-card-title-{index+1}")
        if index:
            line(slide, x - .65, 5.3, x - .10, 5.3, fill=INK, width=1.3, arrow=True, name="fusion-arrow")
    add_concat_icon(slide, 17.55, 4.6)
    add_attention_icon(slide, 20.75, 4.8)
    add_network_icon(slide, 24.55, 4.7)
    add_norm_icon(slide, 28.35, 4.65)


def add_scoring(slide):
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 32.4, 2.45, 5.35, 7.85,
          fill=PALE_GREEN, border=GREEN, width=1.1, dash=True, name="scoring-frame")
    text(slide, "4. Anomaly Scoring", 32.25, 1.45, 5.7, .7, size=14.5, fill=GREEN, bold=True)
    labeled(slide, "Reconstruction\nModule", 33.05, 3.15, 4.05, 2.9,
            fill=PALE_GREEN, border=GREEN, size=11, name="reconstruction-module")
    for index in range(4):
        shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 33.85 + index * .58, 5.1, .42, .52,
              fill="EDF7D9", border=GREEN, width=.7, name="reconstruction-vector")
    line(slide, 35.08, 6.05, 35.08, 6.65, fill=INK, width=1.2, arrow=True)
    labeled(slide, "Anomaly Score", 33.05, 6.7, 4.05, 2.65,
            fill=PALE_GREEN, border=GREEN, size=11, name="anomaly-score")
    # 钟形曲线及红色异常尾部。
    curve = [(33.85, 8.85), (34.25, 8.75), (34.55, 8.25), (34.85, 7.55), (35.15, 7.25), (35.45, 7.55), (35.75, 8.3), (36.1, 8.72)]
    for index, (start, end) in enumerate(zip(curve, curve[1:])):
        line(slide, *start, *end, fill="C62F2F" if index >= 5 else INK, width=1.0, name="anomaly-bell-curve")
    line(slide, 33.75, 8.86, 36.25, 8.86, fill=INK, width=.8)


def add_decision(slide):
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 32.55, 11.75, 5.15, 5.05,
          fill=PALE_PINK, border=PINK, width=1.1, dash=True, name="decision-frame")
    text(slide, "5. Decision", 33.0, 12.05, 4.2, .65, size=14.5, fill=PINK, bold=True)
    labeled(slide, "Thresholding", 33.15, 13.35, 3.95, 1.35,
            fill=WHITE, border=PINK, size=10.5, bold=True, name="thresholding")
    line(slide, 35.12, 14.7, 35.12, 15.15, fill=INK, width=1.1, arrow=True)
    labeled(slide, "Normal / Anomaly", 33.15, 15.2, 3.95, 1.0,
            fill=WHITE, border=PINK, size=10, bold=True, name="normal-anomaly")
    line(slide, 35.12, 10.3, 35.12, 11.72, fill=INK, width=1.5, arrow=True, name="score-to-decision")


def add_training(slide):
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 16.65, 11.55, 9.4, 4.9,
          fill=PALE_BLUE, border=BLUE, width=1.1, dash=True, name="training-frame")
    text(slide, "6. Training Objective", 17.5, 11.75, 7.7, .65, size=14.5, fill=BLUE, bold=True)
    labeled(slide, "Reconstruction Loss\nLrec", 17.05, 12.85, 3.9, 1.95,
            fill=WHITE, border=BLUE, size=10, name="reconstruction-loss")
    text(slide, "+", 21.0, 13.35, .8, .7, size=18)
    labeled(slide, "Regularization Loss\nLreg", 21.7, 12.85, 3.9, 1.95,
            fill=WHITE, border=BLUE, size=10, name="regularization-loss")
    line(slide, 17.4, 15.2, 25.3, 15.2, fill=BLUE, width=.9, name="loss-separator")
    text(slide, "Total Loss: L = Lrec + λLreg", 18.25, 15.35, 6.45, .75, size=11.5)


def add_legend(slide):
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 9.6, 17.8, 20.0, 1.5,
          fill=WHITE, border=GRAY, width=1.0, dash=True, name="legend-frame")
    text(slide, "Legend:", 10.0, 18.1, 1.6, .7, size=10.5, bold=True, align=PP_ALIGN.LEFT)
    line(slide, 13.65, 18.52, 14.85, 18.52, fill=INK, width=1.1, arrow=True, name="legend-data-arrow")
    text(slide, "Data Flow", 15.1, 18.15, 2.2, .65, size=9.5, align=PP_ALIGN.LEFT)
    line(slide, 18.6, 18.52, 19.8, 18.52, fill=INK, width=1.1, dash=True, arrow=True, name="legend-aux-arrow")
    text(slide, "Auxiliary Flow", 20.1, 18.15, 2.8, .65, size=9.5, align=PP_ALIGN.LEFT)
    shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 24.3, 18.25, 1.05, .5, fill=PALE_BLUE, border=BLUE, width=.7, name="legend-vector")
    text(slide, "Vector / Feature", 25.65, 18.15, 3.0, .65, size=9.5, align=PP_ALIGN.LEFT)


def build(output: Path) -> None:
    prs = Presentation(); prs.slide_width = Cm(38.4); prs.slide_height = Cm(20.0)
    while prs.slides:
        rid = prs.slides._sldIdLst[0].rId; prs.part.drop_rel(rid); del prs.slides._sldIdLst[0]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(WHITE)
    add_multivariate_card(slide); add_sequence_encoder(slide); add_pattern_encoder(slide)
    add_fusion(slide); add_scoring(slide); add_decision(slide); add_training(slide); add_legend(slide)
    # 主数据流与辅助训练流向。
    elbow(slide, [(5.6, 7.0), (6.4, 7.0), (6.4, 3.0), (8.3, 3.0)], name="input-to-sequence")
    elbow(slide, [(5.6, 9.0), (6.8, 9.0), (6.8, 10.4), (8.3, 10.4)], name="input-to-pattern")
    elbow(slide, [(14.1, 3.0), (15.35, 3.0), (15.35, 5.25), (16.1, 5.25)], name="sequence-to-fusion")
    elbow(slide, [(14.35, 10.2), (15.15, 10.2), (15.15, 6.8), (16.1, 6.8)], fill=PURPLE, dash=True, name="pattern-to-fusion")
    line(slide, 31.2, 5.3, 32.35, 5.3, fill=INK, width=1.5, arrow=True, name="fusion-to-scoring")
    elbow(slide, [(23.4, 9.45), (23.4, 10.5), (21.35, 10.5), (21.35, 11.5)], name="fusion-to-training")
    elbow(slide, [(16.65, 14.9), (11.7, 14.9), (11.7, 13.35)], fill=BLUE, dash=True, name="training-to-pattern")
    output.parent.mkdir(parents=True, exist_ok=True); prs.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(); parser.add_argument("--input-dir", type=Path, required=True); parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args(); source = args.input_dir / INPUT_NAME
    if not source.is_file() or sha256(source) != INPUT_SHA256:
        raise ValueError("异常检测参考图缺失或 SHA256 不匹配")
    output = args.output_dir / OUTPUT_NAME; build(output); print(output)


if __name__ == "__main__":
    main()
