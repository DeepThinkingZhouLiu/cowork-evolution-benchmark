#!/usr/bin/env python3
"""officeval_054 逐题解法：重建 12.44×7.00 cm 工业园中水回用流程图。"""

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Pt


INPUT_NAME = "路径图.png"
INPUT_SHA256 = "94408780b4c2d64f393a40ea0a222255e87c6c5bcc0e3cdb486bfb059f041ce8"
OUTPUT_NAME = "industrial_water_reuse_editable_process.pptx"
TEAL = "078B81"
DARK_TEAL = "06716B"
NAVY = "0B3E78"
PALE_BAND = "EDF5F8"
PALE_GREEN = "F0F8F3"
PALE_BLUE = "F2F8FC"
GRAY = "5E6770"
WHITE = "FFFFFF"
INK = "111111"
FONT = "SimSun"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def shape(slide, kind, x: float, y: float, w: float, h: float, *,
          fill: str = WHITE, border: str | None = TEAL, width: float = .75,
          dash: bool = False, name: str = ""):
    item = slide.shapes.add_shape(kind, Cm(x), Cm(y), Cm(w), Cm(h))
    item.name = name or f"shape-{len(slide.shapes)}"
    item.fill.solid(); item.fill.fore_color.rgb = rgb(fill)
    if border is None:
        item.line.fill.background()
    else:
        item.line.color.rgb = rgb(border); item.line.width = Pt(width)
        if dash:
            item.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return item


def text(slide, value: str, x: float, y: float, w: float, h: float, *,
         size: float = 7, fill: str = INK, bold: bool = False,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
         font: str = FONT, name: str = ""):
    box = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    box.name = name or f"text-{value[:18]}"
    frame = box.text_frame; frame.clear(); frame.vertical_anchor = anchor
    frame.margin_left = Cm(.01); frame.margin_right = Cm(.01); frame.margin_top = Cm(.01); frame.margin_bottom = Cm(.01)
    paragraph = frame.paragraphs[0]; paragraph.alignment = align; paragraph.space_after = Pt(0); paragraph.line_spacing = 1.0
    run = paragraph.add_run(); run.text = value
    run.font.name = font; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = rgb(fill)
    return box


def line(slide, x1: float, y1: float, x2: float, y2: float, *,
         fill: str = NAVY, width: float = .75, dash: bool = False,
         arrow: bool = False, name: str = ""):
    item = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    item.name = name or f"line-{len(slide.shapes)}"
    item.line.color.rgb = rgb(fill); item.line.width = Pt(width)
    if dash:
        item.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        item.line.end_arrowhead = True
    return item


def add_number(slide, number: int, x: float, y: float):
    shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x, y, .28, .28, fill=TEAL, border=TEAL, name=f"step-number-circle-{number}")
    text(slide, str(number), x, y, .28, .28, size=6.5, fill=WHITE, bold=True,
         font="Times New Roman", name=f"step-number-{number}")


def add_vertical_step(slide, number: int, label: str, x: float, y: float):
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, .86, 2.62,
          fill=WHITE, border=TEAL, width=.75, name=f"step-box-{number}")
    add_number(slide, number, x + .29, y + .16)
    vertical = "\n".join(label) if label != "pH均衡池" else "p\nH\n均\n衡\n池"
    text(slide, vertical, x + .08, y + .53, .70, 1.98, size=7.5, bold=True,
         name=f"step-label-{number}")


def add_aux_icon(slide, kind: str, x: float, y: float):
    if kind == "flask":
        shape(slide, MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE, x, y + .08, .22, .24, fill=PALE_GREEN, border=TEAL, name="aux-flask")
        shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x + .07, y, .08, .12, fill=TEAL, border=TEAL, name="aux-flask-neck")
    elif kind == "drop":
        shape(slide, MSO_AUTO_SHAPE_TYPE.TEAR, x, y, .24, .32, fill=PALE_BLUE, border=TEAL, name="aux-drop")
    elif kind == "shield":
        shape(slide, MSO_AUTO_SHAPE_TYPE.PENTAGON, x, y, .25, .30, fill=PALE_GREEN, border=TEAL, name="aux-shield")
        text(slide, "+", x + .03, y + .03, .19, .20, size=7, fill=TEAL, bold=True)


def add_aux(slide, label: str, x: float, y: float, w: float, kind: str, step_center: float):
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, .42, fill=PALE_GREEN, border=TEAL, name=f"aux-box-{label}")
    add_aux_icon(slide, kind, x + .08, y + .05)
    text(slide, label, x + .35, y + .02, w - .40, .38, size=5.8, fill=DARK_TEAL, name=f"aux-label-{label}")
    line(slide, step_center, y, step_center, y - .42, fill=TEAL, width=.75, dash=True, arrow=True, name=f"aux-link-{label}")


def footer_icon(slide, kind: str, x: float, y: float):
    if kind == "shield":
        shape(slide, MSO_AUTO_SHAPE_TYPE.PENTAGON, x, y, .34, .42, fill=WHITE, border=TEAL, name="footer-shield")
        text(slide, "+", x + .06, y + .08, .22, .20, size=7, fill=TEAL, bold=True)
    elif kind == "clipboard":
        shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y + .04, .32, .38, fill=WHITE, border=TEAL, name="footer-clipboard")
        shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x + .10, y, .12, .09, fill=TEAL, border=TEAL, name="footer-clip")
        for offset in range(2): line(slide, x + .08, y + .17 + offset * .10, x + .25, y + .17 + offset * .10, fill=TEAL, width=.5)
    elif kind == "gear":
        shape(slide, MSO_AUTO_SHAPE_TYPE.GEAR_6, x, y, .40, .40, fill=WHITE, border=TEAL, name="footer-gear")
        shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + .14, y + .14, .12, .12, fill=WHITE, border=TEAL, name="footer-gear-hole")
    else:
        shape(slide, MSO_AUTO_SHAPE_TYPE.FOLDED_CORNER, x, y, .34, .42, fill=WHITE, border=TEAL, name="footer-document")
        for offset in range(3): line(slide, x + .07, y + .14 + offset * .09, x + .25, y + .14 + offset * .09, fill=TEAL, width=.5)


def add_footer_module(slide, x: float, title_value: str, subtitle: str, icon: str):
    footer_icon(slide, icon, x, 11.83)
    text(slide, title_value, x + .48, 11.78, 2.15, .33, size=6, fill=NAVY, bold=True, align=PP_ALIGN.LEFT)
    text(slide, subtitle, x + .48, 12.10, 2.15, .28, size=5, fill=GRAY, align=PP_ALIGN.LEFT)


def build(output: Path) -> None:
    prs = Presentation(); prs.slide_width = Cm(33.867); prs.slide_height = Cm(19.05)
    while prs.slides:
        rid = prs.slides._sldIdLst[0].rId; prs.part.drop_rel(rid); del prs.slides._sldIdLst[0]
    slide = prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(WHITE)

    left, top, total_w, total_h = 10.70, 6.00, 12.44, 7.00
    # 显式的组合底板固定用户指定的 12.44×7.00 cm 总边界。
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, total_w, total_h,
          fill=WHITE, border=None, name="process-group-background")
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, total_w, .80,
          fill=DARK_TEAL, border=DARK_TEAL, name="process-title-background")
    text(slide, "工业园中水回用处理流程与水质保障路径", left + .35, top + .06, total_w - .70, .65,
         size=13, fill=WHITE, bold=True, name="process-title")
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left + .13, top + .78, 12.18, 5.82,
          fill=WHITE, border=TEAL, width=.75, name="process-body-frame")

    step_y = top + 1.68
    step_x0 = left + .42
    gap = .34
    step_w = .86
    labels = ["厂区综合废水收集", "粗格栅与调节井", "pH均衡池", "混凝沉淀池", "多介质过滤器",
              "活性炭吸附塔", "超滤膜组件", "消毒接触池", "回用清水池", "冷却补水系统"]
    centers = []
    for index, label in enumerate(labels):
        x = step_x0 + index * (step_w + gap)
        shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x - .12, top + .85, 1.10, 4.73,
              fill=PALE_BAND if index % 2 == 0 else WHITE, border=None, name=f"process-band-{index+1}")
        add_vertical_step(slide, index + 1, label, x, step_y)
        centers.append(x + step_w / 2)
        if index:
            line(slide, x - gap + .03, step_y + 1.31, x - .03, step_y + 1.31,
                 fill=NAVY, width=.75, arrow=True, name=f"step-arrow-{index}")

    aux_y = top + 4.62
    add_aux(slide, "酸碱调节", left + 2.55, aux_y, 1.25, "flask", centers[2])
    add_aux(slide, "PAC/PAM投加", left + 3.90, aux_y, 1.38, "flask", centers[3])
    add_aux(slide, "核心净化单元", left + 6.55, aux_y, 1.55, "drop", centers[6])
    add_aux(slide, "次氯酸钠投加", left + 8.30, aux_y, 1.62, "shield", centers[7])

    # 第9步向下的回用支路。
    reuse_x, reuse_y = left + 9.55, top + 5.10
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, reuse_x, reuse_y, 2.30, .55,
          fill=WHITE, border=TEAL, name="green-reuse-box")
    leaf = shape(slide, MSO_AUTO_SHAPE_TYPE.TEAR, reuse_x + .17, reuse_y + .10, .25, .32,
                 fill=TEAL, border=TEAL, name="green-reuse-leaf"); leaf.rotation = 45
    text(slide, "绿化与冲洗回用", reuse_x + .47, reuse_y + .06, 1.72, .42, size=6.3, bold=True, name="green-reuse-label")
    line(slide, centers[8], step_y + 2.62, centers[8], reuse_y, fill=NAVY, width=.75, arrow=True, name="clear-pool-down-arrow")

    # 异常回流：沿主体底部左行，再上返第2步。
    return_y = top + 5.42
    line(slide, centers[8] - .15, return_y, centers[1], return_y, fill=NAVY, width=.75, name="abnormal-return-horizontal")
    line(slide, centers[1], return_y, centers[1], step_y + 2.62, fill=NAVY, width=.75, arrow=True, name="abnormal-return-up")
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left + 5.05, top + 5.20, 1.0, .36,
          fill=NAVY, border=NAVY, name="abnormal-return-label-bg")
    text(slide, "异常回流", left + 5.05, top + 5.20, 1.0, .36, size=5.8, fill=WHITE, bold=True, name="abnormal-return-label")

    # 底部保障信息栏。
    footer_y = top + 5.88
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left + .28, footer_y, 11.88, .76,
          fill=PALE_BLUE, border=TEAL, name="assurance-footer")
    for index in range(1, 4):
        xx = left + .28 + index * 2.97
        line(slide, xx, footer_y + .12, xx, footer_y + .64, fill=TEAL, width=.75, dash=True, name="footer-separator")
    add_footer_module(slide, left + .55, "全过程在线监测", "pH/浊度/余氯/电导率", "shield")
    add_footer_module(slide, left + 3.50, "水质达标保障", "满足回用水质标准", "clipboard")
    add_footer_module(slide, left + 6.48, "设备定期维护", "保障系统稳定运行", "gear")
    add_footer_module(slide, left + 9.45, "数据记录追溯", "管理可视化、可追溯", "document")

    output.parent.mkdir(parents=True, exist_ok=True); prs.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(); parser.add_argument("--input-dir", type=Path, required=True); parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args(); source = args.input_dir / INPUT_NAME
    if not source.is_file() or sha256(source) != INPUT_SHA256:
        raise ValueError("工业园中水回用参考图缺失或 SHA256 不匹配")
    output = args.output_dir / OUTPUT_NAME; build(output); print(output)


if __name__ == "__main__":
    main()
