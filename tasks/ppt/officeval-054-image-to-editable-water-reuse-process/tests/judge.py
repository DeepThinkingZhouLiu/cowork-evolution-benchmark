#!/usr/bin/env python3
"""officeval_054 工业园中水回用流程图的独立 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "industrial_water_reuse_editable_process.pptx"
TASK_ID = "cowork-evo/officeval-054-image-to-editable-water-reuse-process"
RUBRIC_VERSION = "omegause-officeval-054-adaptation-v0.1"
JUDGE_VERSION = "officeval-054-adaptation-v0.1"
EMU_PER_CM = 360000.0
STEPS = ["厂区综合废水收集", "粗格栅与调节井", "pH均衡池", "混凝沉淀池", "多介质过滤器", "活性炭吸附塔", "超滤膜组件", "消毒接触池", "回用清水池", "冷却补水系统"]
SPECS = {
    "R001": ("指定 PPTX 可打开，严格为 1 页 16:9", 4, "hurdle"),
    "R002": ("流程图整体为 12.44×7.00 cm 且居中", 5, "hurdle"),
    "R003": ("标题、主体外框与 10 条竖向背景色带完整", 5, "positive"),
    "R004": ("10 个步骤文本、编号、窄圆角框与 9 条水平流向线完整", 7, "hurdle"),
    "R005": ("酸碱、PAC/PAM、核心净化与次氯酸钠四个辅助说明完整", 4, "positive"),
    "R006": ("异常回流与绿化/冲洗回用支路完整", 5, "positive"),
    "R007": ("底部保障栏的 4 个模块、3 条分隔线与图标完整", 5, "positive"),
    "R008": ("文字为宋体，主要边框与线条为 0.75 磅", 3, "positive"),
    "R009": ("页面不含位图，至少含 100 个原生可编辑对象", 5, "hurdle"),
    "R010": ("LibreOffice 可渲染为 1 页，标题、10 个步骤和底部信息可读", 4, "positive"),
    "R011": ("若嵌入位图、缺少两个以上步骤或总尺寸错误，扣 12 分", -12, "penalty"),
}


def cm(value: int) -> float: return value / EMU_PER_CM
def norm(value: str) -> str: return re.sub(r"\s+", "", value).lower()
def count(slide, prefix: str) -> int: return sum(1 for item in slide.shapes if item.name.startswith(prefix))


def criterion(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {"criterion_id": identifier, "score": 1.0 if passed else 0.0, "evidence": evidence, "method": "deterministic",
            "raw": {"description": description, "weight": weight, "type": kind, "passed": bool(passed)}}


def render(path: Path) -> tuple[bool, str]:
    with tempfile.TemporaryDirectory(prefix="officeval054-judge-") as temp_dir:
        completed = subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", temp_dir, str(path)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=120)
        pdf = Path(temp_dir) / f"{path.stem}.pdf"
        if completed.returncode != 0 or not pdf.is_file(): return False, f"转 PDF 失败：{completed.stdout[-300:]}"
        info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, text=True, timeout=30).stdout
        match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE); pages = int(match.group(1)) if match else 0
        txt = Path(temp_dir) / "out.txt"; subprocess.run(["pdftotext", str(pdf), str(txt)], check=True, timeout=30)
        rendered = norm(txt.read_text(errors="replace")); keys = ["工业园中水回用", "消毒接触池", "异常回流", "数据记录追溯"]
        ok = pages == 1 and all(norm(key) in rendered for key in keys)
        return ok, f"页数={pages}，关键文本可读={ok}"


def evaluate(path: Path) -> dict:
    criteria = []
    try: prs = Presentation(path)
    except Exception as exc:
        for identifier in SPECS: criteria.append(criterion(identifier, False, f"打开失败：{exc}"))
        return make_result(path, criteria)
    size_ok = abs(cm(prs.slide_width) - 33.867) <= .05 and abs(cm(prs.slide_height) - 19.05) <= .05
    criteria.append(criterion("R001", len(prs.slides) == 1 and size_ok, f"页数={len(prs.slides)}，页面={cm(prs.slide_width):.3f}×{cm(prs.slide_height):.2f}cm"))
    slide = prs.slides[0] if prs.slides else None
    group = next((item for item in slide.shapes if item.name == "process-group-background"), None) if slide else None
    group_ok = bool(group and abs(cm(group.width) - 12.44) <= .03 and abs(cm(group.height) - 7.00) <= .03 and abs(cm(group.left) - 10.70) <= .05 and abs(cm(group.top) - 6.00) <= .05)
    criteria.append(criterion("R002", group_ok, f"组合坐标/尺寸正确={group_ok}"))
    frame_ok = slide and any(item.name == "process-body-frame" for item in slide.shapes)
    criteria.append(criterion("R003", bool(frame_ok) and count(slide, "process-band-") == 10, f"主体外框={bool(frame_ok)}，色带={count(slide, 'process-band-') if slide else 0}"))
    content = "\n".join(item.text for item in slide.shapes if getattr(item, "has_text_frame", False)) if slide else ""
    missing = [value for value in STEPS if norm(value) not in norm(content)]
    step_counts = [count(slide, prefix) if slide else 0 for prefix in ["step-box-", "step-number-circle-", "step-label-", "step-arrow-"]]
    criteria.append(criterion("R004", not missing and step_counts == [10, 10, 10, 9], f"缺失={missing}，结构={step_counts}"))
    aux_labels = ["酸碱调节", "PAC/PAM投加", "核心净化单元", "次氯酸钠投加"]
    aux_missing = [value for value in aux_labels if norm(value) not in norm(content)]
    criteria.append(criterion("R005", not aux_missing and count(slide, "aux-box-") == 4 and count(slide, "aux-link-") == 4, f"缺失={aux_missing}，说明框/连线={count(slide, 'aux-box-')}/{count(slide, 'aux-link-')}"))
    branch_names = ["abnormal-return-horizontal", "abnormal-return-up", "abnormal-return-label-bg", "green-reuse-box", "green-reuse-leaf", "clear-pool-down-arrow"]
    branch_ok = slide and all(any(item.name == name for item in slide.shapes) for name in branch_names)
    criteria.append(criterion("R006", bool(branch_ok) and all(norm(key) in norm(content) for key in ["异常回流", "绿化与冲洗回用"]), f"支路结构={bool(branch_ok)}"))
    footer_titles = ["全过程在线监测", "水质达标保障", "设备定期维护", "数据记录追溯"]
    footer_ok = all(norm(key) in norm(content) for key in footer_titles) and count(slide, "footer-separator") == 3
    criteria.append(criterion("R007", footer_ok, f"保障标题完整={footer_ok}，分隔线={count(slide, 'footer-separator')}"))
    fonts = []
    widths = []
    if slide:
        for item in slide.shapes:
            if getattr(item, "has_text_frame", False):
                for paragraph in item.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.text.strip() and not item.name.startswith("step-number-"):
                            fonts.append(run.font.name)
            if item.name.startswith(("step-box-", "step-arrow-", "aux-link-", "footer-separator")) and item.line.width:
                widths.append(item.line.width.pt)
    font_ok = fonts and all(value in {"SimSun", "宋体"} for value in fonts)
    width_ok = widths and all(abs(value - .75) <= .03 for value in widths)
    criteria.append(criterion("R008", font_ok and width_ok, f"宋体文本={font_ok}，0.75磅线={width_ok}，检查线数={len(widths)}"))
    pictures = sum(1 for item in slide.shapes if item.shape_type == MSO_SHAPE_TYPE.PICTURE) if slide else 0; natives = len(slide.shapes) - pictures if slide else 0
    criteria.append(criterion("R009", pictures == 0 and natives >= 100, f"图片={pictures}，原生对象={natives}"))
    render_ok, render_evidence = render(path); criteria.append(criterion("R010", render_ok, render_evidence))
    severe = len(prs.slides) != 1 or pictures > 0 or len(missing) >= 2 or not group_ok
    criteria.append(criterion("R011", not severe, f"严重错误={severe}"))
    return make_result(path, criteria)


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [item["criterion_id"] for item in criteria if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0); earned = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0 and item["raw"]["passed"])
    penalties = sum(abs(item["raw"]["weight"]) for item in criteria if item["raw"]["weight"] < 0 and not item["raw"]["passed"]); reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalties) / total))
    return {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [hashlib.sha256(path.read_bytes()).hexdigest()] if path.is_file() else [], "criterion_results": criteria, "reward": reward, "passed": reward >= .8 and not hurdles, "hurdle_failures": hurdles}


def main() -> None:
    parser = argparse.ArgumentParser(); parser.add_argument("--artifact-dir", type=Path, required=True); parser.add_argument("--result", type=Path, required=True); parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args(); output = args.artifact_dir / OUTPUT_NAME
    evaluation = evaluate(output) if output.is_file() else {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [], "criterion_results": [], "reward": 0.0, "passed": False, "hurdle_failures": ["missing_output"]}
    args.result.parent.mkdir(parents=True, exist_ok=True); args.reward_file.parent.mkdir(parents=True, exist_ok=True); args.result.write_text(json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8"); args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8"); print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__": main()
