#!/usr/bin/env python3
"""officeval_055 逐题改造后的 Harbor Judge。

不对照 Golden 文件，而是从上游图片和 Rubric 中拆出可确定检查的
文本、几何、样式、可编辑性和实际渲染条件。同时修正上游 Verifier
的两个已知误判：图例黑箭头不应算主流程箭头，向上的连接线需根据
OOXML ``flipV`` 判断视觉方向。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation

import upstream_verifier


OUTPUT_NAME = "website_outage_editable_flow.pptx"
TASK_ID = "cowork-evo/officeval-055-image-to-editable-outage-flow"
RUBRIC_VERSION = "omegause-officeval-055-adaptation-v0.1"
JUDGE_VERSION = "officeval-055-adaptation-v0.1"
EMU_PER_CM = 360000.0

STAGES = [
    (
        "1",
        "start",
        "At 8:40 a.m., the team noticed that the online store was no longer loading for customers.",
    ),
    (
        "2",
        "check",
        "The support lead checked alerts, confirmed the outage, and opened an incident channel for developers and operations.",
    ),
    (
        "3",
        "problem peak",
        "Orders failed repeatedly because the database connection pool was exhausted during the launch traffic surge.",
    ),
    (
        "4",
        "solution point",
        "The team rolled back the latest configuration and switched traffic to a healthy backup instance.",
    ),
    (
        "5",
        "finish",
        "The site was back up and stable within 10 minutes, order submissions resumed, and a post-incident review has been scheduled for 4:00 PM this afternoon.",
    ),
]

EMOTION_LABELS = ["concerned", "tense and focused", "hopeful", "relieved"]
THEME_TEXT = "theme: managing a website outage before a product launch"

CRITERION_SPECS = {
    "R001": ("指定 PPTX 可打开，严格为 1 页 16:9", 3, "hurdle"),
    "R002": ("五个阶段标题与正文完整、顺序正确，两处指定替换成立", 6, "hurdle"),
    "R003": ("五组编号、图标圆、标题条和正文框布局齐全", 5, "positive"),
    "R004": ("五个阶段图标都由多个可编辑形状组成", 3, "positive"),
    "R005": ("四条主 action 箭头在相邻图标之间水平向右连接", 3, "positive"),
    "R006": ("左上角图例框、action 和 emotion 示例完整", 3, "positive"),
    "R007": ("底部 emotion 折线为六段橙色可编辑虚线，方向与长度正确", 4, "positive"),
    "R008": ("四个情绪标签与底部主题文字顺序、位置和样式正确", 3, "positive"),
    "R009": ("节点圆、标题条、正文框的字体、颜色、尺寸符合来源规格", 4, "positive"),
    "R010": ("整体构图铺满且居中，对象不越出页面", 2, "positive"),
    "R011": ("LibreOffice 实际渲染为 1 页，所有关键文本可读", 3, "positive"),
    "R012": ("页面由可单独选择的文字、形状和连线组成，没有整页压平", 4, "positive"),
    "R013": ("若存在覆盖 60% 以上页面的图片或可编辑对象过少，扣 12 分", -12, "penalty"),
    "R014": ("若缺失两个及以上阶段，或 action / emotion 任一条主线不完整，扣 8 分", -8, "penalty"),
}


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().lower()


def criterion(identifier: str, passed: bool, evidence: str = "") -> dict:
    description, weight, kind = CRITERION_SPECS[identifier]
    return {
        "criterion_id": identifier,
        "description": description,
        "weight": weight,
        "type": kind,
        "passed": bool(passed),
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
    }


def stage_records(model) -> tuple[list[tuple], list[str]]:
    records = []
    errors = []
    for number, title, body in STAGES:
        number_shape = upstream_verifier.number_shape(model, number)
        title_shape = upstream_verifier.stage_title_shape(model, title)
        body_shape = upstream_verifier.stage_body_shape(model, body)
        if number_shape is None:
            errors.append(f"缺少编号 {number}")
        if title_shape is None:
            errors.append(f"缺少标题 {title}")
        if body_shape is None:
            errors.append(f"缺少正文：{body[:32]}")
        records.append((number_shape, title_shape, body_shape))
    return records, errors


def text_and_order_check(model) -> tuple[bool, str, int]:
    records, errors = stage_records(model)
    title_centers = [item[1].center_x_cm for item in records if item[1] is not None]
    if len(title_centers) == 5 and title_centers != sorted(title_centers):
        errors.append(f"阶段标题不是从左到右：{title_centers}")
    if re.search(r"\bassessment\b", normalize(model.all_text)):
        errors.append("第二阶段仍出现 assessment")
    missing_count = sum(1 for record in records if any(value is None for value in record))
    return not errors, "全部文本与顺序正确" if not errors else "；".join(errors), missing_count


def stage_layout_check(model) -> tuple[bool, str]:
    failures = []
    for spec in upstream_verifier.STAGE_SPECS:
        result = upstream_verifier.check_stage(model, spec)
        if not result.passed:
            failures.append(f"{result.rule_id}: {result.evidence}")
    return not failures, "五组阶段节点均符合几何与样式规格" if not failures else "；".join(failures)


def icon_editability_check(model) -> tuple[bool, str]:
    failures = []
    counts = []
    for spec in upstream_verifier.STAGE_SPECS:
        title = upstream_verifier.stage_title_shape(model, spec["title"])
        icons = upstream_verifier.find_icon_candidates(model, title, spec["num"])
        if not icons:
            failures.append(f"阶段 {spec['num']} 缺少可编辑图标圆")
            continue
        icon = min(icons, key=lambda shape: abs(shape.width_cm - 3.2) + abs(shape.height_cm - 3.2))
        local_parts = [
            shape
            for shape in upstream_verifier.nearby_xml_shapes(model, icon.center_x_cm, icon.center_y_cm, 1.55)
            if shape.tag in {"sp", "cxnSp"}
            and shape.name != icon.name
            and 0.01 < shape.area_cm2 < 8.0
        ]
        counts.append(len(local_parts))
        if len(local_parts) < 2:
            failures.append(f"阶段 {spec['num']} 图标内只有 {len(local_parts)} 个可编辑部件")
    return not failures, f"五个图标可编辑部件数={counts}" if not failures else "；".join(failures)


def is_left_to_right_arrow(shape) -> bool:
    if shape.arrow_head == shape.arrow_tail:
        return False
    return (shape.arrow_tail and not shape.flip_h) or (shape.arrow_head and shape.flip_h)


def action_arrow_check(model) -> tuple[bool, str, int]:
    icons = []
    for spec in upstream_verifier.STAGE_SPECS:
        title = upstream_verifier.stage_title_shape(model, spec["title"])
        candidates = upstream_verifier.find_icon_candidates(model, title, spec["num"])
        if candidates:
            icon = min(candidates, key=lambda shape: abs(shape.width_cm - 3.2) + abs(shape.height_cm - 3.2))
            icons.append((icon.center_x_cm, icon.center_y_cm))
    icons.sort()

    arrows = []
    for shape in upstream_verifier.line_candidates(model):
        if not upstream_verifier.is_straight_line_geom(shape):
            continue
        if not (4.0 <= shape.center_y_cm <= 8.0):
            continue
        if abs(shape.height_cm) > 0.15 or not upstream_verifier.color_is(shape.line_color, "black"):
            continue
        if not (2.0 <= shape.length_cm <= 3.0) or not is_left_to_right_arrow(shape) or shape.dashed:
            continue
        arrows.append(shape)

    connected = 0
    if len(icons) == 5:
        for (left_x, left_y), (right_x, _) in zip(icons, icons[1:]):
            if any(left_x < arrow.center_x_cm < right_x and abs(arrow.center_y_cm - left_y) <= 1.5 for arrow in arrows):
                connected += 1
    ok = len(arrows) == 4 and connected == 4
    return ok, f"主流程箭头={len(arrows)}，覆盖相邻间隔={connected}/4", len(arrows)


def legend_check(model) -> tuple[bool, str]:
    result = upstream_verifier.score_legend(model)
    return result.passed, result.evidence


def emotion_line_check(model) -> tuple[bool, str, int]:
    candidates = [
        shape
        for shape in upstream_verifier.emotion_line_candidates(model)
        if upstream_verifier.is_straight_line_geom(shape) and shape.length_cm > 1.0
    ]
    segments = sorted(candidates, key=lambda shape: shape.left_cm)[:6]
    expected_ranges = [(2.5, 3.5), (8.0, 9.5), (6.0, 7.0), (4.5, 6.0), (2.0, 3.0), (6.0, 7.0)]
    lengths = [shape.length_cm for shape in segments]
    length_ok = len(segments) == 6 and all(low <= length <= high for length, (low, high) in zip(lengths, expected_ranges))
    # Connector 的 ext.cy 始终是正数，向上折线通过 flipV 存储。
    directions = ["up" if shape.flip_v else "down" for shape in segments]
    direction_ok = directions == ["down", "up", "down", "up", "down", "up"]
    dash_ok = len(segments) == 6 and all(shape.dashed and 0.8 <= upstream_verifier.effective_line_width_pt(shape.line_width_pt) <= 2.3 for shape in segments)
    arrow_ok = len(segments) == 6 and bool(segments[-1].arrow_head or segments[-1].arrow_tail)
    ok = length_ok and direction_ok and dash_ok and arrow_ok
    evidence = (
        f"段数={len(segments)}，长度={[round(value, 2) for value in lengths]}，"
        f"方向={directions}，虚线={dash_ok}，末端箭头={arrow_ok}"
    )
    return ok, evidence, len(segments)


def labels_and_theme_check(model) -> tuple[bool, str]:
    failures = []
    labels = []
    for label in EMOTION_LABELS:
        matches = upstream_verifier.text_shapes(model, label, exact=True)
        if not matches:
            failures.append(f"缺少 {label}")
            continue
        shape = matches[0]
        labels.append(shape)
        if not upstream_verifier.shape_font_family_ok(shape, ("times", "liberation serif", "cambria", "georgia")):
            failures.append(f"{label} 字体错误")
        if not upstream_verifier.font_size_in_spec(shape, 12, 15):
            failures.append(f"{label} 字号错误")
        if not upstream_verifier.shape_has_text_color(shape, "orange") or not upstream_verifier.shape_italic_ok(shape, True):
            failures.append(f"{label} 颜色或斜体错误")
    if len(labels) == 4 and [shape.center_x_cm for shape in labels] != sorted(shape.center_x_cm for shape in labels):
        failures.append("情绪标签左右顺序错误")
    theme = upstream_verifier.score_theme_text(model)
    if not theme.passed:
        failures.append(theme.evidence)
    return not failures, "四个标签和主题文字正确" if not failures else "；".join(failures)


def composition_check(model) -> tuple[bool, str]:
    bbox = upstream_verifier.main_bbox(model)
    if bbox is None:
        return False, "没有可编辑主内容"
    left, top, right, bottom = bbox
    width_ratio = (right - left) / model.slide_width_cm
    height_ratio = (bottom - top) / model.slide_height_cm
    center_dx = abs((left + right) / 2 - model.slide_width_cm / 2) / model.slide_width_cm
    center_dy = abs((top + bottom) / 2 - model.slide_height_cm / 2) / model.slide_height_cm
    ok = width_ratio >= 0.85 and height_ratio >= 0.75 and center_dx <= 0.06 and center_dy <= 0.08 and upstream_verifier.shapes_inside_slide(model)
    return ok, f"内容包围框={width_ratio:.1%}×{height_ratio:.1%}，中心偏移={center_dx:.1%}/{center_dy:.1%}"


def render_check(path: Path) -> tuple[bool, str]:
    libreoffice = shutil.which("libreoffice")
    pdftotext = shutil.which("pdftotext")
    if not libreoffice or not pdftotext:
        return False, "缺少 libreoffice 或 pdftotext"
    with tempfile.TemporaryDirectory(prefix="cowork-evo-055-") as temp_dir:
        converted = subprocess.run(
            [libreoffice, "--headless", "--convert-to", "pdf", "--outdir", temp_dir, str(path)],
            capture_output=True,
            text=True,
            timeout=180,
            check=False,
        )
        pdf_path = Path(temp_dir) / f"{path.stem}.pdf"
        if converted.returncode != 0 or not pdf_path.is_file():
            return False, (converted.stderr or converted.stdout or "PDF 转换失败")[-800:]
        extracted = subprocess.run(
            [pdftotext, "-layout", str(pdf_path), "-"],
            capture_output=True,
            text=True,
            timeout=60,
            check=False,
        )
        if extracted.returncode != 0:
            return False, (extracted.stderr or "PDF 文本抽取失败")[-800:]
        pages = [page for page in extracted.stdout.split("\f") if page.strip()]
        rendered_text = normalize(" ".join(pages))
    # 五列正文处在同一组水平带上，pdftotext 会把各列按行交错，
    # 不能把整句连续字符串当作渲染结果。每列选取两个唯一短语，
    # 既能确认文本确实渲染，也不会被多列阅读顺序影响。
    required = [title for _, title, _ in STAGES] + EMOTION_LABELS + [THEME_TEXT]
    required.extend(
        [
            "8:40 a.m.",
            "loading for customers",
            "support lead",
            "incident channel",
            "database",
            "connection pool",
            "rolled back",
            "healthy backup",
            "stable within 10 minutes",
            "scheduled for 4:00 PM",
        ]
    )
    missing = [text for text in required if normalize(text) not in rendered_text]
    return len(pages) == 1 and not missing, f"页数={len(pages)}，缺失文本={missing}"


def editability_check(model) -> tuple[bool, str, bool]:
    slide_area = max(model.slide_width_cm * model.slide_height_cm, 1.0)
    large_pictures = [
        shape
        for shape in model.xml_shapes
        if shape.tag == "pic" and shape.area_cm2 / slide_area >= 0.60
    ]
    editable_shapes = model.xml_counts.get("sp", 0)
    connectors = model.xml_counts.get("cxnSp", 0)
    text_shapes = sum(1 for shape in model.shapes if shape.text.strip())
    flattened = bool(large_pictures) or editable_shapes < 25 or connectors < 10 or text_shapes < 16
    evidence = (
        f"大图={len(large_pictures)}，原生形状={editable_shapes}，"
        f"连线={connectors}，文字对象={text_shapes}"
    )
    return not flattened, evidence, flattened


def evaluate(artifact_dir: Path) -> dict:
    output_path = artifact_dir / OUTPUT_NAME
    results = []
    if not output_path.is_file():
        results.append(criterion("R001", False, f"缺少 {OUTPUT_NAME}"))
        return build_result(results, output_path)

    try:
        presentation = Presentation(output_path)
        model = upstream_verifier.inspect_ppt(output_path)
    except Exception as exc:
        results.append(criterion("R001", False, f"文件无法打开：{exc}"))
        return build_result(results, output_path)

    width_cm = presentation.slide_width / EMU_PER_CM
    height_cm = presentation.slide_height / EMU_PER_CM
    file_ok = (
        model.valid_pptx
        and len(presentation.slides) == 1
        and abs(width_cm - 33.867) <= 0.05
        and abs(height_cm - 19.05) <= 0.05
    )
    results.append(criterion("R001", file_ok, f"页数={len(presentation.slides)}，尺寸={width_cm:.3f}×{height_cm:.3f}cm"))
    if not file_ok:
        return build_result(results, output_path)

    text_ok, text_evidence, missing_stages = text_and_order_check(model)
    results.append(criterion("R002", text_ok, text_evidence))

    layout_ok, layout_evidence = stage_layout_check(model)
    results.append(criterion("R003", layout_ok, layout_evidence))

    icons_ok, icons_evidence = icon_editability_check(model)
    results.append(criterion("R004", icons_ok, icons_evidence))

    action_ok, action_evidence, action_count = action_arrow_check(model)
    results.append(criterion("R005", action_ok, action_evidence))

    legend_ok, legend_evidence = legend_check(model)
    results.append(criterion("R006", legend_ok, legend_evidence))

    emotion_ok, emotion_evidence, emotion_count = emotion_line_check(model)
    results.append(criterion("R007", emotion_ok, emotion_evidence))

    labels_ok, labels_evidence = labels_and_theme_check(model)
    results.append(criterion("R008", labels_ok, labels_evidence))

    style_ok = all(upstream_verifier.check_stage(model, spec).passed for spec in upstream_verifier.STAGE_SPECS)
    results.append(criterion("R009", style_ok, "五个节点全部通过来源尺寸、颜色与字体规则" if style_ok else "至少一个节点的视觉样式不合格"))

    composition_ok, composition_evidence = composition_check(model)
    results.append(criterion("R010", composition_ok, composition_evidence))

    rendered, render_evidence = render_check(output_path)
    results.append(criterion("R011", rendered, render_evidence))

    editable, editability_evidence, flattened = editability_check(model)
    results.append(criterion("R012", editable, editability_evidence))

    results.append(criterion("R013", flattened, editability_evidence))
    critical_flow_missing = action_count < 4 or emotion_count < 6
    severe_missing = missing_stages >= 2 or critical_flow_missing
    results.append(
        criterion(
            "R014",
            severe_missing,
            f"缺失阶段={missing_stages}，action={action_count}，emotion={emotion_count}",
        )
    )
    return build_result(results, output_path)


def build_result(results: list[dict], artifact_path: Path) -> dict:
    present = {item["criterion_id"] for item in results}
    for identifier, (_, _, kind) in CRITERION_SPECS.items():
        if identifier not in present:
            results.append(criterion(identifier, False, "因交付文件不存在、损坏或门槛失败，本判据未执行"))
    results.sort(key=lambda item: item["criterion_id"])

    positive_total = sum(item["weight"] for item in results if item["weight"] > 0)
    earned = sum(item["weight"] for item in results if item["weight"] > 0 and item["passed"])
    penalties = sum(item["weight"] for item in results if item["weight"] < 0 and item["passed"])
    reward = max(0.0, min(1.0, (earned + penalties) / positive_total))
    hurdle_failures = [item["criterion_id"] for item in results if item["type"] == "hurdle" and not item["passed"]]
    if hurdle_failures:
        reward = 0.0

    digest = hashlib.sha256(artifact_path.read_bytes() if artifact_path.is_file() else b"").hexdigest()
    normalized_results = []
    for item in results:
        score = (0.0 if item["passed"] else 1.0) if item["type"] == "penalty" else item["score"]
        normalized_results.append(
            {
                "criterion_id": item["criterion_id"],
                "score": score,
                "evidence": item["evidence"],
                "method": item["method"],
                "raw": {
                    "description": item["description"],
                    "weight": item["weight"],
                    "type": item["type"],
                    "passed": item["passed"],
                },
            }
        )
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [digest],
        "criterion_results": normalized_results,
        "reward": round(reward, 6),
        "passed": reward >= 0.8,
        "hurdle_failures": hurdle_failures,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path)
    parser.add_argument("--reward-file", type=Path)
    args = parser.parse_args()
    result = evaluate(args.artifact_dir)
    rendered = json.dumps(result, ensure_ascii=False, indent=2)
    print(rendered)
    if args.result:
        args.result.parent.mkdir(parents=True, exist_ok=True)
        args.result.write_text(rendered + "\n", encoding="utf-8")
    if args.reward_file:
        args.reward_file.parent.mkdir(parents=True, exist_ok=True)
        args.reward_file.write_text(f"{result['reward']:.6f}\n", encoding="utf-8")
    raise SystemExit(0 if result["passed"] else 1)


if __name__ == "__main__":
    main()
