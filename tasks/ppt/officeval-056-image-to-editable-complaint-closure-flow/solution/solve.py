#!/usr/bin/env python3
"""officeval_056 逐题解法：逐对象重建客户投诉闭环处理流程。"""

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Pt


INPUT_NAME = "客户投诉闭环处理流程 (2).png"
INPUT_SHA256 = "a053ad51fac4a0e4a15aa030b3ed3347bbb0a85539e9a6d33d98a9d6c6604217"
OUTPUT_NAME = "customer_complaint_closure_editable.pptx"
FONT = "思源手写体"
NAVY = "0B3A82"
BLUE = "1D66B2"
LIGHT_BLUE = "EDF5FC"
TEAL = "0A9C8F"
LIGHT_GREEN = "EFF9F4"
ORANGE = "F07800"
LIGHT_ORANGE = "FFF6E9"
INK = "202124"
WHITE = "FFFFFF"


def rgb(value: str) -> RGBColor: return RGBColor.from_string(value)
def sha256(path: Path) -> str: return hashlib.sha256(path.read_bytes()).hexdigest()


def shape(slide, kind, x: float, y: float, w: float, h: float, *, fill: str = WHITE,
          border: str | None = NAVY, width: float = 1.0, dash: bool = False, name: str = ""):
    item = slide.shapes.add_shape(kind, Cm(x), Cm(y), Cm(w), Cm(h)); item.name = name or f"shape-{len(slide.shapes)}"
    item.fill.solid(); item.fill.fore_color.rgb = rgb(fill)
    if border is None: item.line.fill.background()
    else:
        item.line.color.rgb = rgb(border); item.line.width = Pt(width)
        if dash: item.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return item


def text(slide, value: str, x: float, y: float, w: float, h: float, *, size: float = 11,
         fill: str = INK, bold: bool = False, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, name: str = ""):
    box = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h)); box.name = name or f"text-{value[:18]}"
    frame = box.text_frame; frame.clear(); frame.vertical_anchor = anchor
    frame.margin_left = Cm(.03); frame.margin_right = Cm(.03); frame.margin_top = Cm(.02); frame.margin_bottom = Cm(.02)
    paragraph = frame.paragraphs[0]; paragraph.alignment = align; paragraph.space_after = Pt(0)
    run = paragraph.add_run(); run.text = value; run.font.name = FONT; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = rgb(fill)
    return box


def line(slide, x1: float, y1: float, x2: float, y2: float, *, fill: str = NAVY,
         width: float = 1.0, dash: bool = False, name: str = ""):
    item = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2)); item.name = name or f"line-{len(slide.shapes)}"
    item.line.color.rgb = rgb(fill); item.line.width = Pt(width)
    if dash: item.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return item


def right_arrow(slide, x: float, y: float, w: float, *, fill: str, name: str):
    return shape(slide, MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW, x, y, w, .38, fill=fill, border=fill, width=.75, name=name)


def down_arrow(slide, x: float, y: float, h: float, *, fill: str, name: str):
    return shape(slide, MSO_AUTO_SHAPE_TYPE.DOWN_ARROW, x, y, .42, h, fill=fill, border=fill, width=.75, name=name)


def icon(slide, kind: str, x: float, y: float, tone: str, prefix: str):
    """每个符号都由多个可编辑几何形状组成。"""
    if kind == "headset":
        shape(slide, MSO_AUTO_SHAPE_TYPE.ARC, x, y, .75, .66, fill=WHITE, border=tone, width=1.2, name=f"{prefix}-headset-arc")
        shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x - .05, y + .35, .18, .36, fill=WHITE, border=tone, name=f"{prefix}-ear-left")
        shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x + .62, y + .35, .18, .36, fill=WHITE, border=tone, name=f"{prefix}-ear-right")
        line(slide, x + .62, y + .68, x + .42, y + .76, fill=tone, name=f"{prefix}-mic")
    elif kind in {"document", "checklist"}:
        shape(slide, MSO_AUTO_SHAPE_TYPE.FOLDED_CORNER, x, y, .72, .82, fill=WHITE, border=tone, name=f"{prefix}-document")
        for idx in range(3):
            line(slide, x + .20, y + .22 + idx * .18, x + .58, y + .22 + idx * .18, fill=tone, width=.7, name=f"{prefix}-doc-line")
            if kind == "checklist": text(slide, "✓", x + .03, y + .12 + idx * .18, .18, .25, size=7, fill=tone, bold=True, name=f"{prefix}-check")
    elif kind == "bars":
        for idx, height in enumerate([.25, .48, .72]):
            shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x + idx * .28, y + .78 - height, .20, height, fill=WHITE, border=tone, name=f"{prefix}-bar")
        line(slide, x - .05, y + .80, x + .82, y + .80, fill=tone, name=f"{prefix}-axis")
    elif kind in {"person", "users"}:
        offsets = [0] if kind == "person" else [-.18, .28]
        for idx, dx in enumerate(offsets):
            shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + .25 + dx, y, .28, .28, fill=WHITE, border=tone, name=f"{prefix}-head-{idx}")
            shape(slide, MSO_AUTO_SHAPE_TYPE.ARC, x + .08 + dx, y + .28, .62, .55, fill=WHITE, border=tone, name=f"{prefix}-body-{idx}")
    elif kind == "folder":
        shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y + .18, .82, .58, fill=WHITE, border=tone, name=f"{prefix}-folder")
        shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x + .08, y + .08, .34, .20, fill=WHITE, border=tone, name=f"{prefix}-folder-tab")
    elif kind == "magnifier":
        shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x, y, .58, .58, fill=WHITE, border=tone, width=1.2, name=f"{prefix}-lens")
        line(slide, x + .45, y + .48, x + .82, y + .85, fill=tone, width=1.3, name=f"{prefix}-handle")
    elif kind == "gear":
        shape(slide, MSO_AUTO_SHAPE_TYPE.GEAR_6, x, y, .85, .85, fill=WHITE, border=tone, width=1.0, name=f"{prefix}-gear")
        shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + .29, y + .29, .27, .27, fill=WHITE, border=tone, name=f"{prefix}-gear-hole")
    elif kind == "shield":
        shape(slide, MSO_AUTO_SHAPE_TYPE.PENTAGON, x, y, .78, .85, fill=WHITE, border=tone, name=f"{prefix}-shield")
        text(slide, "✓", x + .16, y + .18, .45, .35, size=11, fill=tone, bold=True, name=f"{prefix}-shield-check")
    else:  # clipboard
        shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y + .08, .72, .78, fill=WHITE, border=tone, name=f"{prefix}-clipboard")
        shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x + .22, y, .28, .16, fill=WHITE, border=tone, name=f"{prefix}-clip")
        for idx in range(3): line(slide, x + .16, y + .28 + idx * .16, x + .56, y + .28 + idx * .16, fill=tone, width=.7, name=f"{prefix}-clip-line")


def process_box(slide, label: str, kind: str, x: float, y: float, w: float, *,
                tone: str, pale: str, name: str):
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, 2.45, fill=pale, border=tone, width=1.0, name=name)
    icon(slide, kind, x + w / 2 - .42, y + .38, tone, f"{name}-icon")
    text(slide, label, x + .25, y + 1.42, w - .5, .68, size=11, name=f"{name}-label")


def lower_box(slide, label: str, kind: str, x: float, w: float, name: str):
    y = 11.05
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, 2.8, fill=LIGHT_ORANGE, border=ORANGE, width=1.0, name=name)
    icon(slide, kind, x + w / 2 - .42, y + .35, ORANGE, f"{name}-icon")
    text(slide, label, x + .15, y + 1.63, w - .30, .75, size=10.5, name=f"{name}-label")


def build(output: Path) -> None:
    prs = Presentation(); prs.slide_width = Cm(33.867); prs.slide_height = Cm(19.05)
    while prs.slides:
        rid = prs.slides._sldIdLst[0].rId; prs.part.drop_rel(rid); del prs.slides._sldIdLst[0]
    slide = prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb("FCFBF8")

    text(slide, "客户投诉闭环处理流程", 8.1, .55, 17.7, 1.55, size=33, fill=NAVY, bold=True, name="main-title")
    line(slide, 10.85, 2.25, 15.55, 2.25, fill=NAVY, width=2.0, name="title-left-line")
    for idx, tone in enumerate([NAVY, "37A7DD", ORANGE, TEAL]):
        shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 16.05 + idx * .52, 2.10, .22, .22, fill=tone, border=tone, name=f"title-dot-{idx+1}")
    line(slide, 18.28, 2.25, 22.98, 2.25, fill=TEAL, width=2.0, name="title-right-line")

    top_y = 5.0
    boxes = [("top-receive", "接收投诉", "headset", .70, 4.0, NAVY, LIGHT_BLUE),
             ("top-register", "信息登记", "document", 5.50, 4.0, NAVY, LIGHT_BLUE),
             ("top-level", "问题分级", "bars", 10.30, 4.0, NAVY, LIGHT_BLUE),
             ("top-service", "客服处理", "person", 19.50, 3.6, TEAL, LIGHT_GREEN),
             ("top-confirm", "客户确认", "users", 23.90, 3.6, TEAL, LIGHT_GREEN),
             ("top-archive", "结案归档", "folder", 28.30, 4.0, TEAL, LIGHT_GREEN)]
    for name, label, kind, x, w, tone, pale in boxes:
        process_box(slide, label, kind, x, top_y, w, tone=tone, pale=pale, name=name)
    shape(slide, MSO_AUTO_SHAPE_TYPE.DIAMOND, 14.80, 4.55, 3.40, 3.40, fill=WHITE, border=TEAL, width=1.1, name="major-question")
    text(slide, "?", 15.85, 4.82, 1.30, .60, size=17, fill=TEAL, bold=True, name="major-question-icon")
    text(slide, "是否重大\n问题？", 15.18, 5.55, 2.65, 1.45, size=11, name="major-question-label")

    for x1, x2, name in [(4.70, 5.50, "arrow-receive-register"), (9.50, 10.30, "arrow-register-level"),
                         (14.30, 14.80, "arrow-level-question"), (18.20, 19.50, "arrow-question-service"),
                         (23.10, 23.90, "arrow-service-confirm"), (27.50, 28.30, "arrow-confirm-archive")]:
        right_arrow(slide, x1, 6.02, x2 - x1, fill=NAVY, name=name)
    text(slide, "否", 18.35, 5.18, .65, .55, size=15, name="major-no-label")

    # 整改支路。
    lower_box(slide, "指派专员", "person", 7.25, 2.80, "lower-assign")
    lower_box(slide, "原因分析", "magnifier", 10.85, 2.80, "lower-analyze")
    lower_box(slide, "制定整改方案", "checklist", 14.45, 3.20, "lower-plan")
    lower_box(slide, "执行整改", "gear", 18.45, 2.80, "lower-execute")
    lower_box(slide, "质量复核", "shield", 22.05, 2.80, "lower-review")
    shape(slide, MSO_AUTO_SHAPE_TYPE.DIAMOND, 25.65, 10.45, 4.0, 4.0, fill=WHITE, border=ORANGE, width=1.1, name="review-question")
    icon(slide, "clipboard", 27.24, 11.05, ORANGE, "review-question-icon")
    text(slide, "复核通过？", 26.15, 12.15, 3.0, .80, size=10.5, name="review-question-label")
    for x1, x2, name in [(10.05, 10.85, "arrow-assign-analyze"), (13.65, 14.45, "arrow-analyze-plan"),
                         (17.65, 18.45, "arrow-plan-execute"), (21.25, 22.05, "arrow-execute-review"), (24.85, 25.65, "arrow-review-question")]:
        right_arrow(slide, x1, 12.22, x2 - x1, fill=ORANGE, name=name)
    down_arrow(slide, 16.28, 7.75, 3.15, fill=NAVY, name="major-yes-arrow")
    text(slide, "是", 16.72, 8.42, .6, .6, size=15, name="major-yes-label")

    # 复核通过后回到结案归档。
    line(slide, 29.65, 12.45, 31.00, 12.45, fill=NAVY, width=1.1, name="archive-return-right")
    line(slide, 31.00, 12.45, 31.00, 7.90, fill=NAVY, width=1.1, name="archive-return-up")
    shape(slide, MSO_AUTO_SHAPE_TYPE.UP_ARROW, 30.80, 7.25, .40, .75, fill=NAVY, border=NAVY, name="archive-return-tip")
    text(slide, "是", 31.15, 10.45, .6, .6, size=15, name="review-yes-label")
    # 未通过时的虚线返工回路。
    line(slide, 27.65, 14.45, 27.65, 16.85, fill=NAVY, width=1.0, dash=True, name="rework-down")
    line(slide, 27.65, 16.85, 19.85, 16.85, fill=NAVY, width=1.0, dash=True, name="rework-left")
    line(slide, 19.85, 16.85, 19.85, 14.35, fill=NAVY, width=1.0, dash=True, name="rework-up")
    shape(slide, MSO_AUTO_SHAPE_TYPE.UP_ARROW, 19.64, 13.85, .42, .62, fill=NAVY, border=NAVY, name="rework-tip")
    text(slide, "否", 26.90, 14.55, .6, .6, size=15, name="review-no-label")
    text(slide, "返工", 23.60, 16.18, 1.25, .55, size=14, name="rework-label")

    # 左下相关记录说明框。
    shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, .85, 13.15, 5.75, 4.50, fill=WHITE, border=NAVY, width=1.0, name="records-frame")
    text(slide, "相关记录：", 2.55, 13.45, 2.9, .65, size=12, fill=NAVY, bold=True, align=PP_ALIGN.LEFT, name="records-title")
    text(slide, "•  沟通记录\n•  证据材料\n•  处理报告", 2.55, 14.25, 3.1, 2.35, size=11, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name="records-list")
    shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, .65, 14.25, 1.85, 1.85, fill=BLUE, border=NAVY, name="records-icon-bg")
    icon(slide, "clipboard", 1.17, 14.70, WHITE, "records-icon")

    output.parent.mkdir(parents=True, exist_ok=True); prs.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(); parser.add_argument("--input-dir", type=Path, required=True); parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args(); source = args.input_dir / INPUT_NAME
    if not source.is_file() or sha256(source) != INPUT_SHA256: raise ValueError("客户投诉闭环流程参考图缺失或 SHA256 不匹配")
    output = args.output_dir / OUTPUT_NAME; build(output); print(output)


if __name__ == "__main__": main()
