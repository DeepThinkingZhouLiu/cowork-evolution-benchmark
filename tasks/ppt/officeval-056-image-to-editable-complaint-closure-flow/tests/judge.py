#!/usr/bin/env python3
"""officeval_056 客户投诉闭环流程图的独立 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "customer_complaint_closure_editable.pptx"
TASK_ID = "cowork-evo/officeval-056-image-to-editable-complaint-closure-flow"
RUBRIC_VERSION = "omegause-officeval-056-adaptation-v0.1"
JUDGE_VERSION = "officeval-056-adaptation-v0.1"
FONT = "思源手写体"
REQUIRED = ["客户投诉闭环处理流程", "接收投诉", "信息登记", "问题分级", "是否重大问题？", "客服处理", "客户确认", "结案归档", "指派专员", "原因分析", "制定整改方案", "执行整改", "质量复核", "复核通过？", "返工", "相关记录：", "沟通记录", "证据材料", "处理报告"]
SPECS = {
    "R001": ("指定 PPTX 可打开，严格为 1 页 16:9", 4, "hurdle"),
    "R002": ("标题、上方 6 个流程节点、重大问题判断和下方整改链文本完整", 7, "hurdle"),
    "R003": ("上方 6 个圆角流程框、1 个菱形判断与 6 条向右箭头完整", 6, "positive"),
    "R004": ("下方 5 个整改流程框、1 个菱形判断与 5 条橙色箭头完整", 6, "positive"),
    "R005": ("重大问题向下分支、结案归档回流和返工虚线回路完整", 5, "positive"),
    "R006": ("每个业务节点的图标由多个原生形状组成", 4, "positive"),
    "R007": ("左下相关记录框、三个项目和剪贴板图标完整", 4, "positive"),
    "R008": ("所有文字的 OOXML 字体设置为思源手写体", 3, "positive"),
    "R009": ("页面不含位图，且含至少 95 个原生可编辑对象", 5, "hurdle"),
    "R010": ("对象不越界，流程主体铺满页面中部", 3, "positive"),
    "R011": ("LibreOffice 可渲染为 1 页，主要节点和记录可读", 4, "positive"),
    "R012": ("若嵌入位图、缺少两个以上节点或回路不完整，扣 15 分", -15, "penalty"),
}


def norm(value: str) -> str: return re.sub(r"\s+", "", value)
def count(slide, prefix: str) -> int: return sum(1 for item in slide.shapes if item.name.startswith(prefix))


def criterion(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {"criterion_id": identifier, "score": 1.0 if passed else 0.0, "evidence": evidence, "method": "deterministic", "raw": {"description": description, "weight": weight, "type": kind, "passed": bool(passed)}}


def render(path: Path) -> tuple[bool, str]:
    with tempfile.TemporaryDirectory(prefix="officeval056-judge-") as temp_dir:
        completed = subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", temp_dir, str(path)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=120)
        pdf = Path(temp_dir) / f"{path.stem}.pdf"
        if completed.returncode != 0 or not pdf.is_file(): return False, f"转 PDF 失败：{completed.stdout[-300:]}"
        info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, text=True, timeout=30).stdout
        match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE); pages = int(match.group(1)) if match else 0
        txt = Path(temp_dir) / "out.txt"; subprocess.run(["pdftotext", str(pdf), str(txt)], check=True, timeout=30)
        rendered = norm(txt.read_text(errors="replace")); keys = ["客户投诉闭环处理流程", "是否重大问题", "复核通过", "相关记录"]
        ok = pages == 1 and all(norm(key) in rendered for key in keys)
        return ok, f"页数={pages}，关键文本可读={ok}"


def evaluate(path: Path) -> dict:
    criteria = []
    try: prs = Presentation(path)
    except Exception as exc:
        for identifier in SPECS: criteria.append(criterion(identifier, False, f"打开失败：{exc}"))
        return make_result(path, criteria)
    slide = prs.slides[0] if prs.slides else None
    size_ok = abs(prs.slide_width / 360000 - 33.867) <= .05 and abs(prs.slide_height / 360000 - 19.05) <= .05
    criteria.append(criterion("R001", len(prs.slides) == 1 and size_ok, f"页数={len(prs.slides)}，16:9={size_ok}"))
    content = "\n".join(item.text for item in slide.shapes if getattr(item, "has_text_frame", False)) if slide else ""
    missing = [value for value in REQUIRED if norm(value) not in norm(content)]
    criteria.append(criterion("R002", not missing, f"缺失={missing}"))
    top_boxes = sum(1 for name in ["top-receive", "top-register", "top-level", "top-service", "top-confirm", "top-archive"] if slide and any(item.name == name for item in slide.shapes))
    top_arrows = sum(1 for item in slide.shapes if item.name.startswith("arrow-") and any(token in item.name for token in ["receive", "register", "level", "question-service", "service-confirm", "confirm-archive"])) if slide else 0
    criteria.append(criterion("R003", top_boxes == 6 and top_arrows == 6 and slide and any(item.name == "major-question" for item in slide.shapes), f"顶部框={top_boxes}，箭头={top_arrows}"))
    lower_boxes = sum(1 for name in ["lower-assign", "lower-analyze", "lower-plan", "lower-execute", "lower-review"] if slide and any(item.name == name for item in slide.shapes))
    lower_arrows = sum(1 for item in slide.shapes if item.name.startswith("arrow-") and any(token in item.name for token in ["assign", "analyze", "plan-execute", "execute-review", "review-question"])) if slide else 0
    criteria.append(criterion("R004", lower_boxes == 5 and lower_arrows == 5 and slide and any(item.name == "review-question" for item in slide.shapes), f"底部框={lower_boxes}，箭头={lower_arrows}"))
    branch_names = ["major-yes-arrow", "archive-return-right", "archive-return-up", "archive-return-tip", "rework-down", "rework-left", "rework-up", "rework-tip"]
    branch_ok = slide and all(any(item.name == name for item in slide.shapes) for name in branch_names)
    criteria.append(criterion("R005", bool(branch_ok), f"分支/回路结构={bool(branch_ok)}"))
    icon_parts = sum(1 for item in slide.shapes if "-icon-" in item.name or item.name.startswith("records-icon")) if slide else 0
    criteria.append(criterion("R006", icon_parts >= 40, f"可编辑图标部件={icon_parts}"))
    records_ok = slide and all(any(item.name == name for item in slide.shapes) for name in ["records-frame", "records-title", "records-list", "records-icon-bg"])
    criteria.append(criterion("R007", bool(records_ok), f"记录框结构={bool(records_ok)}"))
    fonts = [run.font.name for item in slide.shapes if getattr(item, "has_text_frame", False) for paragraph in item.text_frame.paragraphs for run in paragraph.runs if run.text.strip()] if slide else []
    font_ok = fonts and all(value == FONT for value in fonts)
    criteria.append(criterion("R008", bool(font_ok), f"文字段={len(fonts)}，字体集合={sorted(set(fonts))}"))
    pictures = sum(1 for item in slide.shapes if item.shape_type == MSO_SHAPE_TYPE.PICTURE) if slide else 0; natives = len(slide.shapes) - pictures if slide else 0
    criteria.append(criterion("R009", pictures == 0 and natives >= 95, f"图片={pictures}，原生对象={natives}"))
    inside = slide and all(item.left >= -1000 and item.top >= -1000 and item.left + item.width <= prs.slide_width + 1000 and item.top + item.height <= prs.slide_height + 1000 for item in slide.shapes)
    criteria.append(criterion("R010", bool(inside), f"全部对象在页内={bool(inside)}"))
    render_ok, render_evidence = render(path); criteria.append(criterion("R011", render_ok, render_evidence))
    severe = len(prs.slides) != 1 or pictures > 0 or len(missing) >= 2 or not branch_ok
    criteria.append(criterion("R012", not severe, f"严重错误={severe}"))
    return make_result(path, criteria)


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [item["criterion_id"] for item in criteria if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0); earned = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0 and item["raw"]["passed"]); penalties = sum(abs(item["raw"]["weight"]) for item in criteria if item["raw"]["weight"] < 0 and not item["raw"]["passed"])
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalties) / total))
    return {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [hashlib.sha256(path.read_bytes()).hexdigest()] if path.is_file() else [], "criterion_results": criteria, "reward": reward, "passed": reward >= .8 and not hurdles, "hurdle_failures": hurdles}


def main() -> None:
    parser = argparse.ArgumentParser(); parser.add_argument("--artifact-dir", type=Path, required=True); parser.add_argument("--result", type=Path, required=True); parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args(); output = args.artifact_dir / OUTPUT_NAME
    evaluation = evaluate(output) if output.is_file() else {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [], "criterion_results": [], "reward": 0.0, "passed": False, "hurdle_failures": ["missing_output"]}
    args.result.parent.mkdir(parents=True, exist_ok=True); args.reward_file.parent.mkdir(parents=True, exist_ok=True); args.result.write_text(json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8"); args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8"); print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__": main()
