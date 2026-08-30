#!/usr/bin/env python3
"""officeval_063 的 Harbor Judge。

上游图表规则在第二次遍历 shapes 时使用 Python 包装对象身份比较，
会把图表误判为“与自己重叠 100%”。本 Judge 保留上游文件，并仅将
这一处改为按 OOXML element/shape id 排除图表自身。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import upstream_verifier


OUTPUT_NAME = "直播运营阶段复盘_修订版.pptx"
SOURCE_NAME = "直播运营阶段复盘.pptx"
TASK_ID = "cowork-evo/officeval-063-multi-page-content-editing"
RUBRIC_VERSION = "omegause-officeval-063-adaptation-v0.1"
JUDGE_VERSION = "officeval-063-adaptation-v0.1"
PASS_THRESHOLD = 0.8

CRITERION_SPECS = {
    "R001": ("指定 PPTX 可打开且保持 10 页", 2, "hurdle"),
    "R002": ("未指定页完全保持，指定页的非目标文字与原图保持", 3, "hurdle"),
    "R003": ("第 10 页原生双轴曲线图和 10 个数据点正确", 5, "hurdle"),
    "R004": ("第 2 页第三张岗位卡补齐排班岗标题", 1, "positive"),
    "R005": ("第 5 页卡片、文字和右侧图片调正", 3, "positive"),
    "R006": ("第 5 页背景改为与第 2/3 页接近的浅蓝色", 3, "positive"),
    "R007": ("第 6 页优化动作的重复句去重", 1, "positive"),
    "R008": ("第 9 页只保留近期和长期两个居中板块", 5, "positive"),
    "R009": ("第 9 页长期板块三行指定文字正确", 1, "positive"),
    "R010": ("第 10 页图表标题语义正确", 1, "positive"),
    "R011": ("第 10 页图表尺寸、间距和位置正确", 1, "positive"),
    "R012": ("LibreOffice 实际渲染为 10 页且关键文字可读", 2, "positive"),
    "R013": ("出现圈画截图、批注、临时说明或截图边框", -5, "penalty"),
    "R014": ("页数/页序损坏、修改页被压平或原图丢失", -8, "penalty"),
}


def criterion(identifier: str, passed: bool, evidence: str = "") -> dict:
    description, weight, kind = CRITERION_SPECS[identifier]
    return {
        "criterion_id": identifier,
        "description": description,
        "weight": weight,
        "type": kind,
        "passed": bool(passed),
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
    }


def normalize(text: str) -> str:
    return re.sub(r"\s+", "", text or "")


def text_counter(slide) -> Counter:
    return Counter(
        normalize(shape.text_frame.text)
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and normalize(shape.text_frame.text)
    )


def picture_hashes(slide) -> Counter:
    return Counter(
        hashlib.sha256(shape.image.blob).hexdigest()
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    )


def canonical_slide(slide) -> bytes:
    return etree.tostring(slide._element, method="c14n")


def protection_check(source: Presentation, candidate: Presentation) -> tuple[bool, list[str]]:
    errors = []
    if len(source.slides) != 10 or len(candidate.slides) != 10:
        return False, [f"页数 source={len(source.slides)} candidate={len(candidate.slides)}"]

    for page in (1, 3, 4, 7, 8):
        if canonical_slide(source.slides[page - 1]) != canonical_slide(candidate.slides[page - 1]):
            errors.append(f"p{page} 非目标页 OOXML 发生变化")

    source_p2 = text_counter(source.slides[1])
    candidate_p2 = text_counter(candidate.slides[1])
    expected_p2 = source_p2 + Counter({"排班岗": 1})
    if candidate_p2 != expected_p2:
        errors.append(f"p2 文字差异：缺失={source_p2 - candidate_p2}，新增={candidate_p2 - expected_p2}")
    if picture_hashes(source.slides[1]) != picture_hashes(candidate.slides[1]):
        errors.append("p2 出现额外图片或原图被替换")

    if text_counter(source.slides[4]) != text_counter(candidate.slides[4]):
        errors.append("p5 业务文字发生变化")
    if picture_hashes(source.slides[4]) != picture_hashes(candidate.slides[4]):
        errors.append("p5 右侧原图丢失或被替换")

    source_p6 = source.slides[5]
    candidate_p6 = candidate.slides[5]
    source_other = Counter(
        (shape.name, normalize(shape.text_frame.text))
        for shape in source_p6.shapes
        if getattr(shape, "has_text_frame", False) and shape.name != "Text 8"
    )
    candidate_other = Counter(
        (shape.name, normalize(shape.text_frame.text))
        for shape in candidate_p6.shapes
        if getattr(shape, "has_text_frame", False) and shape.name != "Text 8"
    )
    source_body = next(shape for shape in source_p6.shapes if shape.name == "Text 8")
    candidate_body = next(shape for shape in candidate_p6.shapes if shape.name == "Text 8")
    source_paragraphs = [normalize(paragraph.text) for paragraph in source_body.text_frame.paragraphs if normalize(paragraph.text)]
    expected_paragraphs = []
    seen_paragraphs = Counter()
    for paragraph in source_paragraphs:
        seen_paragraphs[paragraph] += 1
        if paragraph.lstrip("•·-") == normalize("将敏感表达改为场景化、说明式表达") and seen_paragraphs[paragraph] > 1:
            continue
        expected_paragraphs.append(paragraph)
    candidate_paragraphs = [normalize(paragraph.text) for paragraph in candidate_body.text_frame.paragraphs if normalize(paragraph.text)]
    if source_other != candidate_other or candidate_paragraphs != expected_paragraphs:
        errors.append(f"p6 除指定重复句外发生变化：期望段落={expected_paragraphs}，实际={candidate_paragraphs}")
    if picture_hashes(source.slides[5]) != picture_hashes(candidate.slides[5]):
        errors.append("p6 出现额外图片或原图被替换")

    p9_values = text_counter(candidate.slides[8])
    required_p9 = [
        "后续规划",
        "近期",
        "完善话术卡片稳定开播前检查提升互动响应质量",
        "长期",
        "搭建标准化复盘范式强化岗位协同配合效能赋能团队稳步长效发展",
    ]
    missing_p9 = [value for value in required_p9 if p9_values[value] != 1]
    forbidden_p9 = ["中期", "沉淀品类知识库参与活动直播策划协助新人熟悉流程", "稳住基本盘，再把经验变成流程。"]
    present_forbidden = [value for value in forbidden_p9 if p9_values[value]]
    if missing_p9 or present_forbidden:
        errors.append(f"p9 规划文字异常：缺失={missing_p9}，应删未删={present_forbidden}")
    if picture_hashes(source.slides[8]) != picture_hashes(candidate.slides[8]):
        errors.append("p9 背景原图丢失或被替换")

    if text_counter(source.slides[9]) != text_counter(candidate.slides[9]):
        errors.append("p10 原始结束页文字发生变化")
    if picture_hashes(source.slides[9]) != picture_hashes(candidate.slides[9]):
        errors.append("p10 原始照片丢失或被替换")
    chart_count = sum(getattr(shape, "has_chart", False) for shape in candidate.slides[9].shapes)
    if chart_count != 1:
        errors.append(f"p10 原生图表数应为 1，实际 {chart_count}")
    return not errors, errors


def corrected_chart_check(presentation: Presentation) -> tuple[bool, str]:
    slide = presentation.slides[9]
    chart_shapes = [shape for shape in slide.shapes if getattr(shape, "has_chart", False)]
    if len(chart_shapes) != 1:
        return False, f"原生图表数应为 1，实际 {len(chart_shapes)}"
    chart_shape = chart_shapes[0]
    chart = chart_shape.chart
    if "LINE" not in str(chart.chart_type).upper():
        return False, f"图表类型不是折线图：{chart.chart_type}"

    sw, sh = presentation.slide_width, presentation.slide_height
    left, top = chart_shape.left, chart_shape.top
    right, bottom = left + chart_shape.width, top + chart_shape.height
    center_x = left + chart_shape.width / 2
    if not (sw / 3 <= center_x <= sw):
        return False, "图表不在中部或右侧"
    area = max(chart_shape.width * chart_shape.height, 1)
    page_area = max(sw * sh, 1)
    overlaps = []
    for shape in slide.shapes:
        if shape._element is chart_shape._element or shape.shape_id == chart_shape.shape_id:
            continue
        shape_area = shape.width * shape.height
        if shape_area >= page_area * 0.95:
            continue
        if getattr(shape, "has_text_frame", False) and not shape.text_frame.text.strip():
            continue
        inter_x = max(0, min(right, shape.left + shape.width) - max(left, shape.left))
        inter_y = max(0, min(bottom, shape.top + shape.height) - max(top, shape.top))
        if inter_x * inter_y / area > 0.05:
            overlaps.append(shape.name)
    if overlaps:
        return False, f"图表与其他内容重叠：{overlaps}"

    expected_categories = ["1月", "2月", "3月", "4月", "5月"]
    plots = list(chart.plots)
    if len(plots) != 2:
        return False, f"应为双轴两个 plot，实际 {len(plots)}"
    for index, plot in enumerate(plots, start=1):
        categories = [str(value).strip() for value in plot.categories]
        if categories != expected_categories:
            return False, f"plot{index} 月份错误：{categories}"
    series = {item.name: [float(value) for value in item.values] for item in chart.series}
    if set(series) != {"营收数值", "盈利比值"}:
        return False, f"系列名错误：{list(series)}"
    expected_revenue = [68000, 352000, 486000, 423000, 579000]
    expected_ratio = [1.42, 1.56, 1.71, 1.63, 1.85]
    if any(abs(a - b) > 0.5 for a, b in zip(series["营收数值"], expected_revenue)):
        return False, f"营收数据错误：{series['营收数值']}"
    if any(abs(a - b) > 0.001 for a, b in zip(series["盈利比值"], expected_ratio)):
        return False, f"盈利比值错误：{series['盈利比值']}"
    if not chart.has_legend:
        return False, "图表缺少图例"

    chart_xml = chart._chartSpace
    colors = []
    for series_node in chart_xml.findall(".//c:ser", upstream_verifier.NS):
        color = series_node.find(".//c:spPr/a:ln/a:solidFill/a:srgbClr", upstream_verifier.NS)
        colors.append(color.get("val") if color is not None else None)
    if len(colors) != 2 or None in colors or colors[0] == colors[1]:
        return False, f"两条曲线颜色应显式不同：{colors}"
    value_axes = chart_xml.findall(".//c:plotArea/c:valAx", upstream_verifier.NS)
    right_axes = [axis for axis in value_axes if (axis.find("c:axPos", upstream_verifier.NS) is not None and axis.find("c:axPos", upstream_verifier.NS).get("val") == "r")]
    if len(value_axes) != 2 or len(right_axes) != 1:
        return False, f"双数值轴结构错误：valAx={len(value_axes)}，right={len(right_axes)}"
    workbook_part = getattr(chart.part.chart_workbook, "xlsx_part", None)
    if workbook_part is None or not workbook_part.blob:
        return False, "图表缺少嵌入数据工作簿"
    return True, f"双轴原生图表；月份={expected_categories}；营收={series['营收数值']}；盈利={series['盈利比值']}；颜色={colors}"


def upstream_check(function, presentation: Presentation) -> tuple[bool, str]:
    try:
        passed, detail = function(presentation)
        return bool(passed), str(detail)
    except Exception as exc:
        return False, f"上游检查异常：{type(exc).__name__}: {exc}"


def annotation_penalty(path: Path, input_dir: Path, candidate: Presentation) -> tuple[bool, str]:
    annotation_hashes = {
        hashlib.sha256((input_dir / name).read_bytes()).hexdigest()
        for name in ("6.JPG", "7.JPG", "8.JPG", "9.JPG")
    }
    output_picture_hashes = {
        digest
        for slide in candidate.slides
        for digest in picture_hashes(slide)
    }
    leaked = sorted(annotation_hashes & output_picture_hashes)
    suspicious_text = []
    for page, slide in enumerate(candidate.slides, start=1):
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = shape.text_frame.text or ""
            if re.search(r"修订后Task|画圈|红色删掉|黄色|TODO|TBD|待补", text, flags=re.I):
                suspicious_text.append(f"p{page}:{text[:30]}")
    with zipfile.ZipFile(path) as archive:
        comments = [name for name in archive.namelist() if "comment" in name.lower()]
    triggered = bool(leaked or suspicious_text or comments)
    return triggered, f"标注图泄漏={leaked}，临时文字={suspicious_text}，批注部件={comments}"


def render_check(path: Path) -> tuple[bool, str]:
    libreoffice = shutil.which("libreoffice")
    pdfinfo = shutil.which("pdfinfo")
    pdftotext = shutil.which("pdftotext")
    if not libreoffice or not pdfinfo or not pdftotext:
        return False, "缺少 libreoffice、pdfinfo 或 pdftotext"
    with tempfile.TemporaryDirectory(prefix="cowork-evo-063-") as temp_dir:
        converted = subprocess.run(
            [libreoffice, "--headless", "--convert-to", "pdf", "--outdir", temp_dir, str(path)],
            capture_output=True,
            text=True,
            timeout=180,
            check=False,
        )
        pdf_path = Path(temp_dir) / f"{path.stem}.pdf"
        if converted.returncode != 0 or not pdf_path.is_file():
            return False, (converted.stderr or converted.stdout or "PDF 转换失败")[-1000:]
        info = subprocess.run([pdfinfo, str(pdf_path)], capture_output=True, text=True, timeout=60, check=False)
        extracted = subprocess.run([pdftotext, "-layout", str(pdf_path), "-"], capture_output=True, text=True, timeout=60, check=False)
        page_match = re.search(r"^Pages:\s+(\d+)", info.stdout, flags=re.M)
        pages = int(page_match.group(1)) if page_match else 0
        text = normalize(extracted.stdout)
    required = ["排班岗", "将敏感表达改为场景化、说明式表达", "搭建标准化复盘范式", "月度营收与盈利比值趋势", "营收数值", "盈利比值"]
    missing = [value for value in required if normalize(value) not in text]
    return pages == 10 and not missing, f"页数={pages}，缺失文字={missing}"


def evaluate(artifact_dir: Path) -> dict:
    output_path = artifact_dir / OUTPUT_NAME
    results = []
    if not output_path.is_file():
        results.append(criterion("R001", False, f"缺少 {OUTPUT_NAME}"))
        return build_result(results, output_path)
    try:
        candidate = Presentation(output_path)
    except Exception as exc:
        results.append(criterion("R001", False, f"文件无法打开：{exc}"))
        return build_result(results, output_path)

    dim1_ok, dim1_notes = upstream_verifier.check_dim1(str(output_path))
    pages_ok = dim1_ok and len(candidate.slides) == 10
    results.append(criterion("R001", pages_ok, f"上游={dim1_notes}，页数={len(candidate.slides)}"))

    input_dir = Path(os.environ.get("INPUT_DIR", "/app/input_files"))
    source = Presentation(input_dir / SOURCE_NAME)
    protected, protection_errors = protection_check(source, candidate)
    results.append(criterion("R002", protected, "未修改页与非目标内容均保持" if protected else "；".join(protection_errors)))

    chart_ok, chart_evidence = corrected_chart_check(candidate)
    results.append(criterion("R003", chart_ok, chart_evidence))
    mapping = {
        "R004": upstream_verifier.rule_2_paiban,
        "R005": upstream_verifier.rule_5_rotation,
        "R006": upstream_verifier.rule_5_background,
        "R007": upstream_verifier.rule_6_optimize,
        "R008": upstream_verifier.rule_9_two_blocks,
        "R009": upstream_verifier.rule_9_long_term,
        "R010": upstream_verifier.rule_10_chart_title,
        "R011": upstream_verifier.rule_10_chart_size,
    }
    for identifier, function in mapping.items():
        passed, evidence = upstream_check(function, candidate)
        results.append(criterion(identifier, passed, evidence))

    rendered, render_evidence = render_check(output_path)
    results.append(criterion("R012", rendered, render_evidence))
    annotation_hit, annotation_evidence = annotation_penalty(output_path, input_dir, candidate)
    results.append(criterion("R013", annotation_hit, annotation_evidence))
    structural_hit = not pages_ok or not protected
    results.append(
        criterion(
            "R014",
            structural_hit,
            "；".join(protection_errors) if structural_hit else "页数、页序、原图和可编辑对象正常",
        )
    )
    return build_result(results, output_path)


def build_result(results: list[dict], artifact_path: Path) -> dict:
    present = {item["criterion_id"] for item in results}
    for identifier in CRITERION_SPECS:
        if identifier not in present:
            results.append(criterion(identifier, False, "因交付不存在、损坏或门槛失败，本判据未执行"))
    results.sort(key=lambda item: item["criterion_id"])
    positive_total = sum(item["weight"] for item in results if item["weight"] > 0)
    earned = sum(item["weight"] for item in results if item["weight"] > 0 and item["passed"])
    penalties = sum(item["weight"] for item in results if item["weight"] < 0 and item["passed"])
    reward = max(0.0, min(1.0, (earned + penalties) / positive_total))
    hurdle_failures = [item["criterion_id"] for item in results if item["type"] == "hurdle" and not item["passed"]]
    if hurdle_failures:
        reward = 0.0
    digest = hashlib.sha256(artifact_path.read_bytes()).hexdigest() if artifact_path.is_file() else hashlib.sha256(b"").hexdigest()
    normalized = []
    for item in results:
        score = (0.0 if item["passed"] else 1.0) if item["type"] == "penalty" else item["score"]
        normalized.append(
            {
                "criterion_id": item["criterion_id"],
                "score": score,
                "evidence": item["evidence"],
                "method": item["method"],
                "raw": {"description": item["description"], "weight": item["weight"], "type": item["type"], "passed": item["passed"]},
            }
        )
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [digest],
        "criterion_results": normalized,
        "reward": round(reward, 6),
        "passed": reward >= PASS_THRESHOLD,
        "hurdle_failures": hurdle_failures,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    args = parser.parse_args()
    result = evaluate(args.artifact_dir)
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
