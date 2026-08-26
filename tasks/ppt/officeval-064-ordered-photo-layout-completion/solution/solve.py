#!/usr/bin/env python3
"""officeval_064 的逐页参考解：保留模板，按顺序插入并裁剪 12 张真实输入照片。"""

from __future__ import annotations

import argparse
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Cm, Pt


LIGHT_BLUE = RGBColor(173, 216, 230)


def add_center_cropped_picture(slide, image_path: Path, x: float, y: float, w: float, h: float):
    """以输入图片中心为焦点，按原始宽高比裁剪进指定厘米画框。"""
    with Image.open(image_path) as image:
        source_ratio = image.width / image.height
    target_ratio = w / h

    picture = slide.shapes.add_picture(str(image_path), Cm(x), Cm(y), Cm(w), Cm(h))
    if source_ratio > target_ratio:
        crop = (1.0 - target_ratio / source_ratio) / 2.0
        picture.crop_left = crop
        picture.crop_right = crop
    elif source_ratio < target_ratio:
        crop = (1.0 - source_ratio / target_ratio) / 2.0
        picture.crop_top = crop
        picture.crop_bottom = crop
    return picture


def add_title_border(slide):
    """用两个独立、无填充的可编辑矩形组成克制的浅蓝色双线标题边框。"""
    borders = [
        (1.06, 0.32, 19.40, 1.38, 0.75),
        (1.12, 0.38, 19.28, 1.26, 0.40),
    ]
    for x, y, w, h, line_width in borders:
        border = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
        border.fill.background()
        border.line.color.rgb = LIGHT_BLUE
        border.line.width = Pt(line_width)


def build(input_dir: Path, output_path: Path) -> None:
    template = input_dir / "核心技术.pptx"
    presentation = Presentation(template)
    if len(presentation.slides) != 2:
        raise ValueError(f"模板页数应为 2，实际为 {len(presentation.slides)}")

    # 每一项均是人工结合模板内容框、上游尺寸细则和可见文字区域确定的坐标。
    placements = {
        0: [
            (1, 2.49, 10.02, 13.40, 3.00),
            (2, 2.49, 13.20, 6.60, 2.80),
            (3, 9.30, 13.20, 6.55, 2.85),
            (4, 18.03, 10.02, 13.40, 3.00),
            (5, 18.03, 13.20, 6.60, 2.80),
            (6, 24.84, 13.20, 6.55, 2.85),
        ],
        1: [
            (7, 2.49, 10.08, 5.20, 3.00),
            (8, 2.49, 13.26, 5.20, 3.00),
            (9, 7.88, 10.08, 8.00, 6.30),
            (10, 18.03, 10.08, 5.20, 3.00),
            (11, 18.03, 13.26, 5.20, 3.00),
            (12, 23.42, 10.08, 8.00, 6.30),
        ],
    }

    for slide_index, specs in placements.items():
        slide = presentation.slides[slide_index]
        for image_number, x, y, width, height in specs:
            add_center_cropped_picture(
                slide,
                input_dir / f"{image_number}.png",
                x,
                y,
                width,
                height,
            )
        add_title_border(slide)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=Path("/app/input_files"))
    parser.add_argument("--output", type=Path, default=Path("/app/output/核心技术_照片排版版.pptx"))
    args = parser.parse_args()
    build(args.input_dir, args.output)


if __name__ == "__main__":
    main()
