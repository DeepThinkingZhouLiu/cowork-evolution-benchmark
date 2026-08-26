#!/usr/bin/env python3
"""officeval_064 照片顺序、画框几何、裁剪和模板保真的独立 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
import tempfile
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "核心技术_照片排版版.pptx"
TASK_ID = "cowork-evo/officeval-064-ordered-photo-layout-completion"
RUBRIC_VERSION = "omegause-officeval-064-adaptation-v0.1"
JUDGE_VERSION = "officeval-064-adaptation-v0.1"
SPECS = {
    "R001": ("指定 PPTX 可打开，严格为 2 页宽屏", 4, "hurdle"),
    "R002": ("两页各含 6 个独立图片对象，共 12 个", 6, "hurdle"),
    "R003": ("1.png 至 12.png 按四个内容框和视觉位置严格顺序放置", 8, "hurdle"),
    "R004": ("四个内容框内图片尺寸和对齐符合指定版式", 7, "positive"),
    "R005": ("12 张照片均保持原始宽高比并采用对称居中裁剪", 6, "positive"),
    "R006": ("模板原有 46 个对象的文本与几何位置保持不变", 7, "hurdle"),
    "R007": ("两页主标题均有至少两个浅蓝色可编辑形状组成的外围边框", 4, "positive"),
    "R008": ("LibreOffice 可渲染为 2 页，标题、模块名和底部指标可读", 4, "positive"),
    "R009": ("若缺少三张以上图片、顺序错三处以上或原模板严重破坏，扣 18 分", -18, "penalty"),
}

# slide, 图片号, x, y, w, h；逐项来自原 Query 与上游 Rubric，而不是从 Golden 反推。
PLACEMENTS = [
    (0, 1, 2.49, 10.02, 13.40, 3.00), (0, 2, 2.49, 13.20, 6.60, 2.80),
    (0, 3, 9.30, 13.20, 6.55, 2.85), (0, 4, 18.03, 10.02, 13.40, 3.00),
    (0, 5, 18.03, 13.20, 6.60, 2.80), (0, 6, 24.84, 13.20, 6.55, 2.85),
    (1, 7, 2.49, 10.08, 5.20, 3.00), (1, 8, 2.49, 13.26, 5.20, 3.00),
    (1, 9, 7.88, 10.08, 8.00, 6.30), (1, 10, 18.03, 10.08, 5.20, 3.00),
    (1, 11, 18.03, 13.26, 5.20, 3.00), (1, 12, 23.42, 10.08, 8.00, 6.30),
]


def cm(value: int) -> float:
    return value / 360000.0


def criterion(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {
        "criterion_id": identifier,
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
        "raw": {"description": description, "weight": weight, "type": kind, "passed": bool(passed)},
    }


def locate_inputs() -> Path:
    candidates = [Path("/app/input_files"), Path(__file__).resolve().parents[1] / "data/input_files"]
    for candidate in candidates:
        if (candidate / "核心技术.pptx").is_file():
            return candidate
    raise FileNotFoundError("找不到任务输入目录")


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def shape_signature(shape) -> tuple:
    text = shape.text if getattr(shape, "has_text_frame", False) else ""
    return (int(shape.shape_type), shape.left, shape.top, shape.width, shape.height, text)


def render(path: Path) -> tuple[bool, str]:
    with tempfile.TemporaryDirectory(prefix="officeval064-judge-") as temp_dir:
        completed = subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", temp_dir, str(path)],
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            timeout=120,
        )
        pdf = Path(temp_dir) / f"{path.stem}.pdf"
        if completed.returncode != 0 or not pdf.is_file():
            return False, f"转 PDF 失败：{completed.stdout[-300:]}"
        info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, text=True, timeout=30).stdout
        match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
        pages = int(match.group(1)) if match else 0
        text_file = Path(temp_dir) / "text.txt"
        subprocess.run(["pdftotext", str(pdf), str(text_file)], check=True, timeout=30)
        rendered = re.sub(r"\s+", "", text_file.read_text(errors="replace"))
        required = ["高效动力链方案", "智能变速总成", "高强复合半轴", "整车性能协同优化", "尾部导流组件", "座舱轻量布局", "18%", "15%"]
        ok = pages == 2 and all(re.sub(r"\s+", "", value) in rendered for value in required)
        return ok, f"页数={pages}，关键文本可读={ok}"


def border_ok(slide, title_text: str) -> tuple[bool, int]:
    title = next((shape for shape in slide.shapes if getattr(shape, "has_text_frame", False) and title_text in shape.text), None)
    if title is None:
        return False, 0
    tx1, ty1, tx2, ty2 = cm(title.left), cm(title.top), cm(title.left + title.width), cm(title.top + title.height)
    candidates = []
    for shape in slide.shapes:
        if shape is title or shape.shape_type not in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM):
            continue
        if getattr(shape, "has_text_frame", False) and shape.text.strip():
            continue
        x1, y1, x2, y2 = cm(shape.left), cm(shape.top), cm(shape.left + shape.width), cm(shape.top + shape.height)
        contains = x1 <= tx1 + .02 and y1 <= ty1 + .02 and x2 >= tx2 - .02 and y2 >= ty2 - .02
        try:
            color = shape.line.color.rgb
            light_blue = color is not None and color[2] > color[0] and color[2] > color[1] and min(color) >= 120
        except Exception:
            light_blue = False
        if contains and light_blue:
            candidates.append(shape)
    return len(candidates) >= 2, len(candidates)


def evaluate(path: Path) -> dict:
    criteria = []
    try:
        inputs = locate_inputs()
        source = Presentation(inputs / "核心技术.pptx")
        submission = Presentation(path)
    except Exception as exc:
        for identifier in SPECS:
            criteria.append(criterion(identifier, False, f"打开失败：{exc}"))
        return make_result(path, criteria)

    page_ok = len(submission.slides) == 2 and abs(cm(submission.slide_width) - 33.865) <= .05 and abs(cm(submission.slide_height) - 19.05) <= .05
    criteria.append(criterion("R001", page_ok, f"页数={len(submission.slides)}，页面={cm(submission.slide_width):.3f}×{cm(submission.slide_height):.2f}cm"))

    pictures_by_slide = [[shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE] for slide in submission.slides]
    counts = [len(items) for items in pictures_by_slide]
    picture_count_ok = counts == [6, 6]
    criteria.append(criterion("R002", picture_count_ok, f"逐页图片数={counts}"))

    expected_hashes = {number: sha256((inputs / f"{number}.png").read_bytes()) for number in range(1, 13)}
    matched = []
    order_errors = []
    geometry_errors = []
    crop_errors = []
    for slide_index, number, x, y, width, height in PLACEMENTS:
        pictures = pictures_by_slide[slide_index] if slide_index < len(pictures_by_slide) else []
        if not pictures:
            order_errors.append(number)
            geometry_errors.append(number)
            crop_errors.append(number)
            continue
        target_cx, target_cy = x + width / 2, y + height / 2
        available = [item for item in pictures if id(item) not in {id(value) for value in matched}]
        if not available:
            order_errors.append(number)
            geometry_errors.append(number)
            crop_errors.append(number)
            continue
        picture = min(available, key=lambda item: abs(cm(item.left + item.width / 2) - target_cx) + abs(cm(item.top + item.height / 2) - target_cy))
        matched.append(picture)
        if sha256(picture.image.blob) != expected_hashes[number]:
            order_errors.append(number)
        actual = (cm(picture.left), cm(picture.top), cm(picture.width), cm(picture.height))
        expected = (x, y, width, height)
        if any(abs(a - b) > .12 for a, b in zip(actual, expected)):
            geometry_errors.append(number)
        with Image.open(inputs / f"{number}.png") as image:
            source_ratio = image.width / image.height
        visible_width = 1.0 - picture.crop_left - picture.crop_right
        visible_height = 1.0 - picture.crop_top - picture.crop_bottom
        visible_ratio = source_ratio * visible_width / visible_height if visible_height > 0 else 0
        symmetric = abs(picture.crop_left - picture.crop_right) <= .005 and abs(picture.crop_top - picture.crop_bottom) <= .005
        if abs(visible_ratio - width / height) > .02 or not symmetric:
            crop_errors.append(number)

    criteria.append(criterion("R003", not order_errors, f"顺序/来源错误图片号={order_errors}"))
    criteria.append(criterion("R004", not geometry_errors, f"位置或尺寸错误图片号={geometry_errors}"))
    criteria.append(criterion("R005", not crop_errors, f"宽高比或对称裁剪错误图片号={crop_errors}"))

    preservation_errors = []
    if len(source.slides) != len(submission.slides):
        preservation_errors.append("页数")
    else:
        for slide_index, (source_slide, output_slide) in enumerate(zip(source.slides, submission.slides), 1):
            source_signatures = [shape_signature(shape) for shape in source_slide.shapes]
            output_signatures = [shape_signature(shape) for shape in list(output_slide.shapes)[:len(source_signatures)]]
            if source_signatures != output_signatures:
                preservation_errors.append(f"第{slide_index}页原对象")
    with zipfile.ZipFile(path) as archive:
        has_comments = any(name.startswith("ppt/comments/") or name.startswith("ppt/commentAuthors") for name in archive.namelist())
    if has_comments:
        preservation_errors.append("批注")
    criteria.append(criterion("R006", not preservation_errors, f"模板保真错误={preservation_errors}"))

    border_results = []
    for slide, title in zip(submission.slides, ["高效动力链方案", "整车性能协同优化"]):
        ok, count = border_ok(slide, title)
        border_results.append((ok, count))
    criteria.append(criterion("R007", all(ok for ok, _ in border_results), f"逐页浅蓝边框构件={border_results}"))

    render_passed, render_evidence = render(path)
    criteria.append(criterion("R008", render_passed, render_evidence))

    severe = sum(abs(6 - count) for count in counts) >= 3 or len(order_errors) >= 3 or bool(preservation_errors)
    criteria.append(criterion("R009", not severe, f"严重错误={severe}"))
    return make_result(path, criteria)


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [item["criterion_id"] for item in criteria if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0 and item["raw"]["passed"])
    penalties = sum(abs(item["raw"]["weight"]) for item in criteria if item["raw"]["weight"] < 0 and not item["raw"]["passed"])
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalties) / total))
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [sha256(path.read_bytes())] if path.is_file() else [],
        "criterion_results": criteria,
        "reward": reward,
        "passed": reward >= .8 and not hurdles,
        "hurdle_failures": hurdles,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    output = args.artifact_dir / OUTPUT_NAME
    evaluation = evaluate(output) if output.is_file() else {
        "task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION,
        "artifact_sha256": [], "criterion_results": [], "reward": 0.0, "passed": False,
        "hurdle_failures": ["missing_output"],
    }
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
