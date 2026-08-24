#!/usr/bin/env python3
"""officeval_066 的逐题参考解法。

本文件只处理这道真实任务的两个原始输入：读取 Excel 中现有的 10×8 数据表，
在原始四页 PPT 后追加一页可编辑表格和一页原生可编辑折线图。它不是任务生成器，
也不创造 Query、输入数据或答案事实。
"""

from __future__ import annotations

import argparse
import re
from copy import deepcopy
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


INPUT_XLSX = "PPT相关不同数据表.xlsx"
INPUT_PPTX = "表格图表关联数据.pptx"
OUTPUT_PPTX = "表格图表关联数据_新增表格折线图.pptx"

TITLE_TABLE = "门店数字化与低碳运营指标"
TITLE_CHART = "门店数字化与低碳运营指标趋势"
BACKGROUND = "F7FAFC"
HEADER_FILL = "2A9D8F"
ROW_FILLS = ("FFFFFF", "EDF4F7")
TEXT_DARK = "25364A"
TEXT_MUTED = "66788A"
LINE_COLORS = ("2F66A8", "D9901A", "6C4AA8", "2A9D8F", "D65A6F", "5B7C99", "8C6D31")
MARKERS = (
    XL_MARKER_STYLE.CIRCLE,
    XL_MARKER_STYLE.SQUARE,
    XL_MARKER_STYLE.DIAMOND,
    XL_MARKER_STYLE.TRIANGLE,
    XL_MARKER_STYLE.PLUS,
    XL_MARKER_STYLE.STAR,
    XL_MARKER_STYLE.X,
)


def read_source_table(path: Path) -> list[list[object]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    if workbook.sheetnames != ["关联不同数据"]:
        raise ValueError(f"预期唯一工作表‘关联不同数据’，实际 {workbook.sheetnames}")
    sheet = workbook[workbook.sheetnames[0]]
    values = [list(row) for row in sheet.iter_rows(values_only=True)]
    while values and not any(value is not None for value in values[-1]):
        values.pop()
    if len(values) != 10 or any(len(row) != 8 for row in values):
        raise ValueError(f"预期 10×8 数据，实际 {len(values)}×{max(map(len, values), default=0)}")
    expected_years = list(range(2016, 2025))
    years = [int(row[0]) for row in values[1:]]
    if years != expected_years:
        raise ValueError(f"年份应为 2016–2024，实际 {years}")
    for row in values[1:]:
        if any(not isinstance(value, (int, float)) for value in row[1:]):
            raise ValueError(f"发现非数值指标行：{row}")
    return values


def set_background(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(BACKGROUND)


def copy_template_decorations(source_slide, target_slide) -> None:
    """复制原页左上和右下两个无外部关系的装饰平行四边形。"""
    decorations = [shape for shape in source_slide.shapes if shape.name.startswith("Parallelogram")]
    if len(decorations) != 2:
        raise ValueError(f"预期 2 个模板装饰对象，实际 {len(decorations)}")
    for shape in decorations:
        target_slide.shapes._spTree.insert_element_before(deepcopy(shape.element), "p:extLst")


def add_title(slide, text: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.42), Inches(7.4), Inches(0.55))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(TEXT_DARK)


def set_cell_text(cell, value: object, *, header: bool, first_column: bool = False) -> None:
    frame = cell.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = Inches(0.025)
    frame.margin_right = Inches(0.025)
    frame.margin_top = Inches(0.015)
    frame.margin_bottom = Inches(0.015)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    if isinstance(value, float):
        run.text = f"{value:.1f}"
    else:
        run.text = str(value)
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(8.5 if header else 10.5)
    run.font.bold = header or first_column
    run.font.color.rgb = RGBColor.from_string("FFFFFF" if header else TEXT_DARK)


def add_table_slide(presentation: Presentation, rows: list[list[object]]) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    set_background(slide)
    copy_template_decorations(presentation.slides[2], slide)
    add_title(slide, TITLE_TABLE)

    table_shape = slide.shapes.add_table(10, 8, Inches(0.65), Inches(1.18), Inches(12.0), Inches(5.18))
    table = table_shape.table
    table.columns[0].width = Inches(0.74)
    metric_width = (Inches(12.0) - Inches(0.74)) // 7
    for column_index in range(1, len(table.columns)):
        table.columns[column_index].width = metric_width
    table.rows[0].height = Inches(0.72)
    for row_index in range(1, len(table.rows)):
        table.rows[row_index].height = Inches(0.495)

    for row_index, row_values in enumerate(rows):
        for column_index, value in enumerate(row_values):
            cell = table.cell(row_index, column_index)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(
                HEADER_FILL if row_index == 0 else ROW_FILLS[(row_index - 1) % 2]
            )
            set_cell_text(
                cell,
                value,
                header=row_index == 0,
                first_column=column_index == 0 and row_index > 0,
            )

    note = slide.shapes.add_textbox(Inches(0.69), Inches(6.51), Inches(5.3), Inches(0.24))
    frame = note.text_frame
    frame.clear()
    run = frame.paragraphs[0].add_run()
    run.text = "数据来源：PPT相关不同数据表.xlsx；单位见各指标表头。"
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(TEXT_MUTED)


def short_metric_name(value: object) -> str:
    return re.sub(r"（.*?）", "", str(value)).strip()


def add_chart_slide(presentation: Presentation, rows: list[list[object]]) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    set_background(slide)
    copy_template_decorations(presentation.slides[3], slide)
    add_title(slide, TITLE_CHART)

    chart_data = ChartData()
    chart_data.categories = [str(int(row[0])) for row in rows[1:]]
    for column_index, header in enumerate(rows[0][1:], start=1):
        chart_data.add_series(short_metric_name(header), [float(row[column_index]) for row in rows[1:]])

    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        Inches(0.72),
        Inches(1.16),
        Inches(11.92),
        Inches(4.92),
        chart_data,
    )
    chart = chart_shape.chart
    chart.has_title = False
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.legend.font.name = "Microsoft YaHei"
    chart.legend.font.size = Pt(9)

    value_axis = chart.value_axis
    value_axis.minimum_scale = 0
    value_axis.maximum_scale = 100
    value_axis.major_unit = 20
    value_axis.has_major_gridlines = True
    value_axis.tick_labels.font.name = "Microsoft YaHei"
    value_axis.tick_labels.font.size = Pt(9)
    category_axis = chart.category_axis
    category_axis.tick_labels.font.name = "Microsoft YaHei"
    category_axis.tick_labels.font.size = Pt(9)

    for index, series in enumerate(chart.series):
        color = RGBColor.from_string(LINE_COLORS[index])
        series.format.line.color.rgb = color
        series.format.line.width = Pt(2.0)
        series.marker.style = MARKERS[index]
        series.marker.size = 6
        series.marker.format.fill.solid()
        series.marker.format.fill.fore_color.rgb = color
        series.marker.format.line.color.rgb = color

    highest = max(
        (float(value), int(row[0]), short_metric_name(rows[0][column_index]))
        for row in rows[1:]
        for column_index, value in enumerate(row[1:], start=1)
    )
    note = slide.shapes.add_textbox(Inches(0.76), Inches(6.18), Inches(6.7), Inches(0.28))
    frame = note.text_frame
    frame.clear()
    run = frame.paragraphs[0].add_run()
    run.text = f"最高点：{highest[2]}在{highest[1]}年达到 {highest[0]:.1f}"
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(TEXT_MUTED)


def solve(input_dir: Path, output_dir: Path) -> Path:
    rows = read_source_table(input_dir / INPUT_XLSX)
    presentation = Presentation(input_dir / INPUT_PPTX)
    if len(presentation.slides) != 4:
        raise ValueError(f"预期原 PPT 为 4 页，实际 {len(presentation.slides)} 页")
    add_table_slide(presentation, rows)
    add_chart_slide(presentation, rows)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / OUTPUT_PPTX
    presentation.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=Path("/app/data/input_files"))
    parser.add_argument("--output-dir", type=Path, default=Path("/app/output"))
    args = parser.parse_args()
    solve(args.input_dir, args.output_dir)


if __name__ == "__main__":
    main()
