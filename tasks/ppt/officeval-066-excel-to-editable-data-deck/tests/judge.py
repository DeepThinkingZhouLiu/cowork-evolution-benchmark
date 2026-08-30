#!/usr/bin/env python3
"""officeval_066 的 Harbor Judge。

直接复用上游逐题 Verifier，并额外核对原四页保留性、原生可编辑对象、
大面积截图覆盖和 LibreOffice 实际渲染结果。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import upstream_verifier


OUTPUT_NAME = "表格图表关联数据_新增表格折线图.pptx"
INPUT_NAME = "表格图表关联数据.pptx"
TASK_ID = "cowork-evo/officeval-066-excel-to-editable-data-deck"
RUBRIC_VERSION = "omegause-officeval-066-adaptation-v0.1"
JUDGE_VERSION = "officeval-066-adaptation-v0.1"
PASS_THRESHOLD = 0.8

CRITERION_SPECS = {
    "R001": ("指定 PPTX 可打开，且页序为原四页、表格页、折线图页", 2, "hurdle"),
    "R002": ("原四页保持不变，新增表格和图表均为原生可编辑对象", 2, "hurdle"),
    "R003": ("第 5 页 10×8 表格及全部真实数据准确", 5, "positive"),
    "R004": ("第 5 页标题文字、位置和字体正确", 1, "positive"),
    "R005": ("第 5 页表格配色与交替行样式正确", 1, "positive"),
    "R006": ("第 6 页原生折线图、嵌入工作簿及 63 个数据点正确", 5, "positive"),
    "R007": ("第 6 页七条折线颜色与数据点标记清晰", 1, "positive"),
    "R008": ("第 6 页标题和真实最高点说明正确", 1, "positive"),
    "R009": ("新增两页背景与原四页一致且为 F7FAFC", 3, "positive"),
    "R010": ("LibreOffice 实际渲染为 6 页且关键文字可读", 2, "positive"),
    "R011": ("新增内容被截图压平、缺少原生对象或页序严重损坏", -8, "penalty"),
    "R012": ("前四页多处被改动，或出现批注、标记与多余对象", -6, "penalty"),
}


def criterion(identifier: str, passed: bool, evidence: str = "") -> dict:
    description, weight, kind = CRITERION_SPECS[identifier]
    return {
        "criterion_id": identifier,
        "description": description,
        "weight": weight,
        "type": kind,
        "passed": bool(passed),
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
    }


def shape_signature(shape) -> dict:
    signature = {
        "name": shape.name,
        "type": int(shape.shape_type),
        "geometry": [shape.left, shape.top, shape.width, shape.height],
        "text": shape.text_frame.text if getattr(shape, "has_text_frame", False) else None,
        "table": None,
        "chart": None,
    }
    if getattr(shape, "has_table", False):
        table = shape.table
        signature["table"] = [
            [table.cell(row, column).text for column in range(len(table.columns))]
            for row in range(len(table.rows))
        ]
    if getattr(shape, "has_chart", False):
        chart = shape.chart
        plots = []
        for plot in chart.plots:
            plots.append(
                {
                    "categories": [str(value) for value in plot.categories],
                    "series": [
                        {
                            "name": series.name,
                            "values": [None if value is None else round(float(value), 8) for value in series.values],
                        }
                        for series in plot.series
                    ],
                }
            )
        signature["chart"] = {"type": str(chart.chart_type), "plots": plots}
    return signature


def slide_signature(slide) -> dict:
    return {
        "background": upstream_verifier.get_slide_bg_hex(slide),
        "shapes": [shape_signature(shape) for shape in slide.shapes],
    }


def compare_preserved_slides(source: Presentation, candidate: Presentation) -> list[str]:
    errors = []
    if len(source.slides) != 4:
        return [f"基准输入页数不是 4，而是 {len(source.slides)}"]
    if len(candidate.slides) < 4:
        return [f"交付只有 {len(candidate.slides)} 页，无法保留前 4 页"]
    for index in range(4):
        before = slide_signature(source.slides[index])
        after = slide_signature(candidate.slides[index])
        if before != after:
            before_shapes = before["shapes"]
            after_shapes = after["shapes"]
            if len(before_shapes) != len(after_shapes):
                errors.append(f"第 {index + 1} 页对象数 {len(before_shapes)} -> {len(after_shapes)}")
            else:
                changed = [
                    position
                    for position, (left, right) in enumerate(zip(before_shapes, after_shapes), start=1)
                    if left != right
                ]
                errors.append(f"第 {index + 1} 页对象发生变化：{changed[:20]}")
            if before["background"] != after["background"]:
                errors.append(
                    f"第 {index + 1} 页背景 {before['background']} -> {after['background']}"
                )
    return errors


def native_editability(candidate: Presentation) -> tuple[bool, list[str], bool]:
    errors = []
    flattened = False
    if len(candidate.slides) != 6:
        return False, [f"页数应为 6，实际 {len(candidate.slides)}"], True
    table_shapes = [shape for shape in candidate.slides[4].shapes if getattr(shape, "has_table", False)]
    chart_shapes = [shape for shape in candidate.slides[5].shapes if getattr(shape, "has_chart", False)]
    if len(table_shapes) != 1:
        errors.append(f"第 5 页原生表格数应为 1，实际 {len(table_shapes)}")
    elif len(table_shapes[0].table.rows) != 10 or len(table_shapes[0].table.columns) != 8:
        errors.append("第 5 页原生表格不是 10×8")
    if len(chart_shapes) != 1:
        errors.append(f"第 6 页原生图表数应为 1，实际 {len(chart_shapes)}")
    elif "LINE" not in str(chart_shapes[0].chart.chart_type).upper():
        errors.append(f"第 6 页图表不是折线图：{chart_shapes[0].chart.chart_type}")

    page_area = candidate.slide_width * candidate.slide_height
    for slide_index in (4, 5):
        for shape_index, shape in enumerate(candidate.slides[slide_index].shapes, start=1):
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            ratio = shape.width * shape.height / page_area
            if ratio >= 0.60:
                flattened = True
                errors.append(f"第 {slide_index + 1} 页图片{shape_index}覆盖 {ratio:.0%} 页面")
    if not table_shapes or not chart_shapes:
        flattened = True
    return not errors, errors, flattened


def upstream_check(function, candidate: Presentation) -> tuple[bool, str]:
    try:
        passed, detail = function(candidate)
        return bool(passed), str(detail)
    except Exception as exc:
        return False, f"上游检查异常：{type(exc).__name__}: {exc}"


def render_check(path: Path) -> tuple[bool, str]:
    libreoffice = shutil.which("libreoffice")
    pdftotext = shutil.which("pdftotext")
    if not libreoffice or not pdftotext:
        return False, "缺少 libreoffice 或 pdftotext"
    with tempfile.TemporaryDirectory(prefix="cowork-evo-066-") as temp_dir:
        converted = subprocess.run(
            [libreoffice, "--headless", "--convert-to", "pdf", "--outdir", temp_dir, str(path)],
            capture_output=True,
            text=True,
            timeout=180,
            check=False,
        )
        pdf_path = Path(temp_dir) / f"{path.stem}.pdf"
        if converted.returncode != 0 or not pdf_path.is_file():
            return False, (converted.stderr or converted.stdout or "PDF 转换失败")[-1000:]
        extracted = subprocess.run(
            [pdftotext, "-layout", str(pdf_path), "-"],
            capture_output=True,
            text=True,
            timeout=60,
            check=False,
        )
        if extracted.returncode != 0:
            return False, (extracted.stderr or "PDF 文字抽取失败")[-1000:]
        pages = [page for page in extracted.stdout.split("\f") if page.strip()]
        normalized = re.sub(r"\s+", "", "".join(pages))
    required = [
        "门店数字化与低碳运营指标",
        "门店数字化与低碳运营指标趋势",
        "2016",
        "2024",
        "智慧门店响应率",
        "运营协同指数",
        "最高点",
        "78.9",
    ]
    missing = [text for text in required if re.sub(r"\s+", "", text) not in normalized]
    return len(pages) == 6 and not missing, f"页数={len(pages)}，缺失文字={missing}"


def evaluate(artifact_dir: Path) -> dict:
    output_path = artifact_dir / OUTPUT_NAME
    results = []
    if not output_path.is_file():
        results.append(criterion("R001", False, f"缺少 {OUTPUT_NAME}"))
        return build_result(results, output_path)
    try:
        candidate = Presentation(output_path)
    except Exception as exc:
        results.append(criterion("R001", False, f"文件无法打开：{exc}"))
        return build_result(results, output_path)

    pages_ok, pages_detail = upstream_check(upstream_verifier.check_pages_order, candidate)
    results.append(criterion("R001", pages_ok, pages_detail))

    source_path = Path(os.environ.get("INPUT_DIR", "/app/input_files")) / INPUT_NAME
    preservation_errors = []
    try:
        source = Presentation(source_path)
        preservation_errors = compare_preserved_slides(source, candidate)
    except Exception as exc:
        preservation_errors = [f"无法比较基准输入：{exc}"]
    editable, edit_errors, flattened = native_editability(candidate)
    integrity_ok = not preservation_errors and editable
    integrity_evidence = (
        "前 4 页语义与对象保持不变；新增表格和折线图均为原生可编辑对象"
        if integrity_ok
        else "；".join((preservation_errors + edit_errors)[:30])
    )
    results.append(criterion("R002", integrity_ok, integrity_evidence))

    check_map = {
        "R003": upstream_verifier.check_s5_table_structure,
        "R004": upstream_verifier.check_s5_title,
        "R005": upstream_verifier.check_s5_table_style,
        "R006": upstream_verifier.check_s6_chart_data,
        "R007": upstream_verifier.check_s6_line_style,
        "R008": upstream_verifier.check_s6_title,
        "R009": upstream_verifier.check_bg_color,
    }
    for identifier, function in check_map.items():
        passed, evidence = upstream_check(function, candidate)
        results.append(criterion(identifier, passed, evidence))

    rendered, render_evidence = render_check(output_path)
    results.append(criterion("R010", rendered, render_evidence))
    results.append(
        criterion(
            "R011",
            flattened or not pages_ok,
            "；".join(edit_errors) if flattened or not pages_ok else "未发现截图压平、原生对象缺失或页序损坏",
        )
    )
    deductions = upstream_verifier.check_deductions(candidate, str(output_path))
    preservation_penalty = len(preservation_errors) >= 2
    results.append(
        criterion(
            "R012",
            preservation_penalty or bool(deductions),
            f"保留性错误={preservation_errors}；上游扣分项={deductions}",
        )
    )
    return build_result(results, output_path)


def build_result(results: list[dict], artifact_path: Path) -> dict:
    present = {item["criterion_id"] for item in results}
    for identifier in CRITERION_SPECS:
        if identifier not in present:
            results.append(criterion(identifier, False, "因交付不存在、损坏或门槛失败，本判据未执行"))
    results.sort(key=lambda item: item["criterion_id"])

    positive_total = sum(item["weight"] for item in results if item["weight"] > 0)
    earned = sum(item["weight"] for item in results if item["weight"] > 0 and item["passed"])
    penalties = sum(item["weight"] for item in results if item["weight"] < 0 and item["passed"])
    reward = max(0.0, min(1.0, (earned + penalties) / positive_total))
    hurdle_failures = [
        item["criterion_id"]
        for item in results
        if item["type"] == "hurdle" and not item["passed"]
    ]
    if hurdle_failures:
        reward = 0.0

    digest = (
        hashlib.sha256(artifact_path.read_bytes()).hexdigest()
        if artifact_path.is_file()
        else hashlib.sha256(b"").hexdigest()
    )
    normalized = []
    for item in results:
        score = (0.0 if item["passed"] else 1.0) if item["type"] == "penalty" else item["score"]
        normalized.append(
            {
                "criterion_id": item["criterion_id"],
                "score": score,
                "evidence": item["evidence"],
                "method": item["method"],
                "raw": {
                    "description": item["description"],
                    "weight": item["weight"],
                    "type": item["type"],
                    "passed": item["passed"],
                },
            }
        )
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [digest],
        "criterion_results": normalized,
        "reward": round(reward, 6),
        "passed": reward >= PASS_THRESHOLD,
        "hurdle_failures": hurdle_failures,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    args = parser.parse_args()
    result = evaluate(args.artifact_dir)
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
