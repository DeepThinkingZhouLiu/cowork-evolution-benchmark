#!/usr/bin/env python3
"""逐题重放 officeval_068 的参考交付物。

这不是通用制题脚本。代码只处理本题两份固定真实输入：人工重建 PDF 的
10 页为可编辑对象，保留并重排 PPT 的 15 页，然后统一为浅蓝视觉系统。
"""

from __future__ import annotations

import argparse
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUTPUT_NAME = "合并_浅蓝可编辑版.pptx"
SOURCE_PPT = "综合调度坐席系统操作介绍.pptx"

FONT = "Microsoft YaHei"
BG = "F2F8FF"
NAVY = "173B63"
BODY = "314860"
MUTED = "6E8298"
BLUE = "2B6DEB"
TEAL = "149E9E"
GREEN = "35A779"
ORANGE = "E9912A"
WHITE = "FFFFFF"
PALE_BLUE = "E7F1FF"
PALE_TEAL = "E5F5FA"
PALE_ORANGE = "FFF2E0"
BORDER = "B8CDE4"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_slide_background(slide) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BG)


def set_shape_fill(shape, color: str) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)


def set_shape_line(shape, color: str = BORDER, width: float = 1.0) -> None:
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 10,
    color: str = BODY,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
    margin: float = 0.03,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.line_spacing = 1.0
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return shape


def add_rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: str = WHITE,
    line: str = BORDER,
    rounded: bool = True,
):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    set_shape_fill(shape, fill)
    set_shape_line(shape, line)
    return shape


def add_circle(slide, x: float, y: float, diameter: float, *, fill: str = PALE_BLUE, line: str = BLUE):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(diameter), Inches(diameter)
    )
    set_shape_fill(shape, fill)
    set_shape_line(shape, line)
    return shape


def add_chevron(slide, x: float, y: float, w: float = 0.24, h: float = 0.34, color: str = BLUE):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    set_shape_fill(shape, color)
    shape.line.fill.background()
    return shape


def add_divider(slide, x: float, y: float, w: float, color: str = BORDER):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.015)
    )
    set_shape_fill(shape, color)
    shape.line.fill.background()
    return shape


def add_footer(slide, page: int, brand: str = "应急通信保障 · 调度坐席协同") -> None:
    add_divider(slide, 0.65, 6.98, 12.0, "D4E2F0")
    add_text(slide, brand, 0.65, 7.02, 4.6, 0.25, size=8, color=MUTED)
    add_text(slide, str(page), 12.25, 7.02, 0.42, 0.25, size=8, color=MUTED, align=PP_ALIGN.RIGHT)


def add_title(slide, title: str, page: int, section: str, subtitle: str = "") -> None:
    set_slide_background(slide)
    add_text(slide, section, 0.65, 0.32, 3.0, 0.26, size=8, color=TEAL, bold=True)
    add_text(slide, title, 0.65, 0.62, 11.8, 0.53, size=26, color=NAVY, bold=True)
    add_divider(slide, 0.65, 1.27, 12.0)
    if subtitle:
        add_text(slide, subtitle, 0.68, 1.34, 11.7, 0.34, size=10, color=MUTED)
    add_footer(slide, page)


def add_card(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    title: str,
    body: str,
    *,
    accent: str = BLUE,
    badge: str | None = None,
):
    add_rect(slide, x, y, w, h, fill=WHITE, line=BORDER)
    strip = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h)
    )
    set_shape_fill(strip, accent)
    strip.line.fill.background()
    title_x = x + 0.28
    if badge is not None:
        add_circle(slide, x + 0.22, y + 0.22, 0.38, fill=PALE_BLUE, line=accent)
        add_text(slide, badge, x + 0.22, y + 0.22, 0.38, 0.38, size=10, color=accent, bold=True, align=PP_ALIGN.CENTER)
        title_x = x + 0.72
    add_text(slide, title, title_x, y + 0.17, w - (title_x - x) - 0.18, 0.34, size=10, color=NAVY, bold=True)
    add_text(slide, body, x + 0.24, y + 0.62, w - 0.46, h - 0.78, size=10, color=BODY, valign=MSO_ANCHOR.TOP)


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def build_pdf_page_1(prs) -> None:
    slide = blank_slide(prs)
    set_slide_background(slide)
    add_circle(slide, 10.42, 0.52, 2.2, fill=PALE_TEAL, line=TEAL)
    add_chevron(slide, 11.18, 1.25, 0.55, 0.72, TEAL)
    add_text(slide, "专题汇报", 0.82, 1.18, 2.2, 0.30, size=10, color=TEAL, bold=True)
    add_text(slide, "应急有线人工通信保障网", 0.78, 1.65, 8.9, 0.68, size=26, color=NAVY, bold=True)
    add_text(slide, "定义、场景、价值与台站体系解析", 0.80, 2.48, 6.8, 0.38, size=10, color=MUTED)
    add_rect(slide, 0.80, 3.22, 7.35, 1.52, fill=WHITE)
    add_text(slide, "保留人工判断与手动路由能力，在自动平台受限或现场条件复杂时，为关键岗位提供稳定、可控、可追踪的有线语音联络。", 1.12, 3.62, 6.75, 0.76, size=10, color=BODY, valign=MSO_ANCHOR.TOP)
    for index, (label, accent) in enumerate((("人工接续", BLUE), ("有线传输", TEAL), ("层级汇接", GREEN), ("闭环记录", ORANGE))):
        x = 0.82 + index * 1.80
        add_rect(slide, x, 5.25, 1.55, 0.48, fill=WHITE, line=BORDER)
        add_circle(slide, x + 0.12, 5.37, 0.22, fill=PALE_BLUE, line=accent)
        add_text(slide, label, x + 0.42, 5.30, 1.02, 0.36, size=10, color=NAVY, bold=True)
    add_footer(slide, 1, "应急有线人工通信保障网")


def build_pdf_page_2(prs) -> None:
    slide = blank_slide(prs)
    add_title(slide, "目录", 2, "内容导航", "从网络定位到设备与运行闭环")
    items = [
        ("01", "网络定位与构成", "人工接续、有线链路、用户终端与安全记录", BLUE),
        ("02", "典型保障场景", "常态值守、突发联络、现场保障与跨班协同", TEAL),
        ("03", "核心能力价值", "稳定、可控、优先级调度、低依赖与快速恢复", ORANGE),
        ("04", "台站层级与设备", "中心汇接、区域转接、现场保障与终端接入", GREEN),
    ]
    for row, (number, title, body, accent) in enumerate(items):
        y = 1.83 + row * 1.18
        add_card(slide, 1.10, y, 11.15, 0.92, title, body, accent=accent, badge=number)


def build_pdf_page_3(prs) -> None:
    slide = blank_slide(prs)
    add_title(slide, "网络定位与构成", 3, "01 · 网络定位", "面向应急场景的人工有线语音保障链路")
    add_card(
        slide, 0.70, 1.86, 4.55, 3.95, "什么是人工通信保障网？",
        "由人工坐席进行呼叫受理、线路选择与跨站接续的有线语音网络。它不追求完全自动化，而是保留人工判断、身份核对和手动路由能力，在自动平台受限或现场条件复杂时提供稳定联络通道。",
        accent=ORANGE,
    )
    components = [
        ("人工接续台", "坐席受理、路由选择、通话保持与转接", BLUE),
        ("有线传输链路", "专线、电缆、配线架等物理通道", TEAL),
        ("用户终端", "固定电话、现场话机、值守终端", GREEN),
        ("安全与记录", "身份核对、权限控制、操作日志留存", ORANGE),
    ]
    for index, (title, body, accent) in enumerate(components):
        x = 5.62 + (index % 2) * 3.35
        y = 1.86 + (index // 2) * 1.96
        add_card(slide, x, y, 3.05, 1.62, title, body, accent=accent, badge=str(index + 1))
    add_text(slide, "关键词：人工接续 · 有线传输 · 层级汇接 · 闭环记录", 0.90, 6.13, 7.4, 0.34, size=10, color=TEAL, bold=True)


def build_pdf_page_4(prs) -> None:
    slide = blank_slide(prs)
    add_title(slide, "典型保障场景", 4, "02 · 使用场景", "从常态值守到突发联络的备用语音链路")
    scenarios = [
        ("应急指挥备份", "自动平台拥塞或中断时，保持关键岗位之间的语音联络。", BLUE),
        ("园区专线联络", "用于调度中心、门岗、机房与保障小组之间的固定通道。", TEAL),
        ("偏远站点接入", "在覆盖不足区域，通过有线端点形成基础通话能力。", ORANGE),
        ("跨班值守交接", "将未完成事项、回呼任务和线路状态交接给下一班组。", GREEN),
        ("大型活动保障", "临时设置接续坐席，对外联络与内部协同分层处理。", BLUE),
        ("灾害抢险协同", "现场通道受限时，保留低依赖、可快速恢复的联络方式。", TEAL),
    ]
    for index, (title, body, accent) in enumerate(scenarios):
        x = 0.72 + (index % 3) * 4.18
        y = 1.82 + (index // 3) * 2.15
        add_card(slide, x, y, 3.78, 1.75, title, body, accent=accent, badge=f"{index + 1:02d}")


def build_pdf_page_5(prs) -> None:
    slide = blank_slide(prs)
    add_title(slide, "核心能力价值（一）", 5, "03 · 核心价值", "强调可用、可控、可审计")
    add_card(
        slide, 0.86, 1.98, 11.58, 1.66, "稳定性",
        "以物理线路和人工接续为基础，减少对复杂平台、外部网络和单点系统的依赖；当主用通道受限时，仍能维持基础语音联络。",
        accent=ORANGE, badge="01",
    )
    add_card(
        slide, 0.86, 4.05, 11.58, 1.66, "可控性",
        "坐席在接续前进行对象确认、权限判断和路由选择；关键通话过程可按岗位规则记录，方便值守复核与责任闭环。",
        accent=TEAL, badge="02",
    )


def build_pdf_page_6(prs) -> None:
    slide = blank_slide(prs)
    add_title(slide, "核心能力价值（二）", 6, "03 · 核心价值", "让应急联络有顺序、有冗余、有恢复路径")
    values = [
        ("优先级调度", "按岗位、事件等级和线路状态安排接续顺序，保障关键联络先行。", ORANGE),
        ("快速恢复", "断点定位、备用通道切换和人工确认流程简单，利于现场抢修。", TEAL),
        ("低依赖运行", "不强依赖复杂自动化平台，适合电力、网络或设备条件受限环境。", GREEN),
        ("易培训维护", "操作规则清晰，设备构成直观，值守人员可在短周期内掌握。", BLUE),
    ]
    for index, (title, body, accent) in enumerate(values):
        x = 0.85 + (index % 2) * 6.03
        y = 1.92 + (index // 2) * 2.14
        add_card(slide, x, y, 5.58, 1.72, title, body, accent=accent, badge=f"{index + 1:02d}")


def build_pdf_page_7(prs) -> None:
    slide = blank_slide(prs)
    add_title(slide, "台站体系与协作架构", 7, "04 · 台站层级", "层级汇接、分区管理、按需接入")
    nodes = [
        ("中心汇接站", "跨区域调度与总控协调，维护核心线路资源。", BLUE),
        ("区域转接站", "承接片区联络与邻区互通，分担话务压力。", TEAL),
        ("现场保障站", "靠近任务现场，负责临时通道和人员协同。", ORANGE),
        ("终端接入点", "固定岗亭、值守室或移动保障单元的接入端。", GREEN),
    ]
    for index, (title, body, accent) in enumerate(nodes):
        x = 0.67 + index * 3.12
        add_card(slide, x, 2.05, 2.68, 2.38, title, body, accent=accent, badge=str(index + 1))
        if index < 3:
            add_chevron(slide, x + 2.78, 3.02, 0.24, 0.36, accent)
    add_rect(slide, 1.02, 5.10, 11.28, 0.88, fill=PALE_TEAL, line=BORDER)
    add_text(slide, "通过层级化管理，网络可在日常值守、突发调度和跨区域协同之间切换，既保留人工判断，又避免所有通话集中到单一节点。", 1.30, 5.28, 10.72, 0.48, size=10, color=BODY)


def build_pdf_page_8(prs) -> None:
    slide = blank_slide(prs)
    add_title(slide, "典型设备组成", 8, "04 · 台站设备", "用自绘可编辑图形替代实拍设备，保持表达一致")
    equipment = [
        ("人工交换坐席", "受理来话、填写记录、选择线路、保持或转接通话。", ORANGE),
        ("线路分配架", "将多组线缆按区域、用途和优先级集中管理。", TEAL),
        ("有线传输介质", "包含专线、电缆和临时敷设线缆，承担语音传输。", BLUE),
        ("安全记录设备", "记录接续时间、对象、动作和异常，便于复核。", GREEN),
        ("值守终端", "承载联系人检索、回呼提醒、交接清单和日志查询。", TEAL),
    ]
    for index, (title, body, accent) in enumerate(equipment[:3]):
        add_card(slide, 0.70 + index * 4.18, 1.87, 3.78, 1.72, title, body, accent=accent, badge=f"{index + 1:02d}")
    for index, (title, body, accent) in enumerate(equipment[3:], start=3):
        add_card(slide, 1.76 + (index - 3) * 5.24, 4.03, 4.78, 1.72, title, body, accent=accent, badge=f"{index + 1:02d}")


def build_pdf_page_24(prs) -> None:
    slide = blank_slide(prs)
    add_title(slide, "运行流程与管理要点", 24, "运行闭环", "让人工链路形成可追踪的闭环")
    steps = [
        ("接入确认", "识别来话来源与联络诉求"),
        ("记录需求", "填写对象、号码和处理备注"),
        ("选择通道", "查看线路状态并确定路由"),
        ("呼出转接", "建立双方或多方通话"),
        ("归档复核", "保存记录并处理回呼任务"),
    ]
    for index, (title, body) in enumerate(steps):
        x = 0.55 + index * 2.55
        add_card(slide, x, 1.85, 2.12, 1.56, title, body, accent=(BLUE, TEAL, ORANGE, GREEN, BLUE)[index], badge=str(index + 1))
        if index < 4:
            add_chevron(slide, x + 2.18, 2.43, 0.22, 0.34, (BLUE, TEAL, ORANGE, GREEN)[index])
    add_text(slide, "三条值守原则", 0.72, 4.02, 3.0, 0.34, size=10, color=TEAL, bold=True)
    principles = [
        ("可靠接入", "先确认通路与语音焦点，再执行后续动作。"),
        ("身份核对", "接续前核对岗位、部门或约定口令。"),
        ("闭环记录", "每次呼入、转接、回呼和异常均形成日志。"),
    ]
    for index, (title, body) in enumerate(principles):
        add_card(slide, 0.72 + index * 4.18, 4.52, 3.76, 1.28, title, body, accent=(BLUE, TEAL, ORANGE)[index])


def build_pdf_page_25(prs) -> None:
    slide = blank_slide(prs)
    set_slide_background(slide)
    add_circle(slide, 9.72, 0.72, 2.35, fill=PALE_BLUE, line=BLUE)
    add_chevron(slide, 10.58, 1.56, 0.58, 0.78, TEAL)
    add_text(slide, "谢谢聆听", 0.92, 1.70, 7.2, 0.62, size=26, color=NAVY, bold=True)
    add_text(slide, "THANKS FOR LISTENING", 0.94, 2.46, 4.8, 0.32, size=10, color=TEAL, bold=True)
    add_rect(slide, 0.92, 3.30, 9.25, 1.35, fill=WHITE, line=BORDER)
    add_text(slide, "稳定 · 可控 · 可恢复", 1.24, 3.58, 4.0, 0.32, size=10, color=NAVY, bold=True)
    add_text(slide, "人工有线通信链路为应急联络保留可靠的组织保障。", 1.24, 4.02, 7.9, 0.34, size=10, color=BODY)
    add_text(slide, "通信保障研究组", 0.94, 5.54, 3.2, 0.32, size=10, color=MUTED)
    add_footer(slide, 25, "应急通信保障 · 完整交付")


def clone_source_slide(prs: Presentation, source_slide):
    slide = blank_slide(prs)
    set_slide_background(slide)
    for shape in source_slide.shapes:
        element = deepcopy(shape.element)
        slide.shapes._spTree.insert_element_before(element, "p:extLst")
    return slide


def shape_max_font_size(shape) -> float:
    if not getattr(shape, "has_text_frame", False):
        return 0.0
    sizes = []
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size is not None:
                sizes.append(run.font.size.pt)
    return max(sizes, default=0.0)


def normalize_cloned_slide(slide, page: int, source_index: int) -> None:
    """保留源页对象与文字，只统一背景、字体、标题层级和连续页码。"""

    set_slide_background(slide)
    text_shapes = [
        shape for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip()
    ]
    title_candidates = [
        shape for shape in text_shapes
        if not shape.text_frame.text.strip().isdigit()
        and not (shape.top > Inches(6.8) and shape.left > Inches(11.5))
    ]
    title_shape = max(title_candidates, key=shape_max_font_size) if title_candidates else None

    for shape in text_shapes:
        raw_text = shape.text_frame.text.strip()
        is_page_number = (
            raw_text.isdigit()
            and shape.top > Inches(6.8)
            and shape.left > Inches(11.5)
        )
        if is_page_number:
            shape.text_frame.clear()
            paragraph = shape.text_frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.RIGHT
            run = paragraph.add_run()
            run.text = str(page)
            run.font.name = FONT
            run.font.size = Pt(8)
            run.font.color.rgb = rgb(MUTED)
            continue

        for paragraph in shape.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.LEFT
            paragraph.line_spacing = 1.0
            for run in paragraph.runs:
                run.font.name = FONT
                run.font.size = Pt(26 if shape is title_shape else 10)
                run.font.bold = True if shape is title_shape else run.font.bold
                run.font.color.rgb = rgb(NAVY if shape is title_shape else BODY)

    for shape in slide.shapes:
        try:
            is_rounded = shape.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE
        except ValueError:
            is_rounded = False
        if is_rounded:
            set_shape_fill(shape, WHITE)
            set_shape_line(shape, BORDER, 1.0)
            try:
                shape.adjustments[0] = 0.16667
            except (IndexError, ValueError):
                pass

    if source_index == 1:
        # 原封面含深蓝整页矩形和深色控制台示意；改成浅蓝背景与白色面板。
        for shape in slide.shapes:
            if shape.width >= Inches(13.0) and shape.height >= Inches(7.2):
                set_shape_fill(shape, BG)
                set_shape_line(shape, BG)
                continue
            try:
                fill_rgb = str(shape.fill.fore_color.rgb) if shape.fill.type == 1 and shape.fill.fore_color.rgb else ""
            except Exception:
                fill_rgb = ""
            if fill_rgb in {"172A4A", "20385F", "2E4B7B"}:
                set_shape_fill(shape, WHITE if fill_rgb == "20385F" else PALE_BLUE)
                set_shape_line(shape, BORDER)

    # 若源页没有可识别的右下页码，则补一个，不改动源内容。
    numbers = [
        shape.text_frame.text.strip()
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
        and shape.top > Inches(6.8)
        and shape.left > Inches(11.5)
    ]
    if str(page) not in numbers:
        add_text(slide, str(page), 12.25, 7.02, 0.42, 0.25, size=8, color=MUTED, align=PP_ALIGN.RIGHT)


def build_deck(input_dir: Path, output_dir: Path) -> Path:
    source = Presentation(input_dir / SOURCE_PPT)
    if len(source.slides) != 18:
        raise ValueError(f"源 PPT 应为 18 页，实际 {len(source.slides)} 页")

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)

    build_pdf_page_1(prs)
    build_pdf_page_2(prs)
    build_pdf_page_3(prs)
    build_pdf_page_4(prs)
    build_pdf_page_5(prs)
    build_pdf_page_6(prs)
    build_pdf_page_7(prs)
    build_pdf_page_8(prs)

    # 删除源 PPT 原第 13、16、18 页；其余 15 页按原顺序插入 PDF 第 8 页之后。
    keep_source_pages = list(range(1, 13)) + [14, 15, 17]
    for page, source_page in enumerate(keep_source_pages, start=9):
        slide = clone_source_slide(prs, source.slides[source_page - 1])
        normalize_cloned_slide(slide, page, source_page)

    build_pdf_page_24(prs)
    build_pdf_page_25(prs)
    if len(prs.slides) != 25:
        raise AssertionError(f"最终应为 25 页，实际 {len(prs.slides)} 页")

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / OUTPUT_NAME
    prs.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    output = build_deck(args.input_dir, args.output_dir)
    print(output)


if __name__ == "__main__":
    main()
