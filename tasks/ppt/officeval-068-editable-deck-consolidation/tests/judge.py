#!/usr/bin/env python3
"""officeval_068 的 Harbor Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import subprocess
import tempfile
from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import upstream_verifier


OUTPUT_NAME = "合并_浅蓝可编辑版.pptx"
SOURCE_PPT = "综合调度坐席系统操作介绍.pptx"
TASK_ID = "cowork-evo/officeval-068-editable-deck-consolidation"
RUBRIC_VERSION = "omegause-officeval-068-adaptation-v0.1"
JUDGE_VERSION = "officeval-068-adaptation-v0.1"
PASS_THRESHOLD = 0.8

CRITERION_SPECS = {
    "R001": ("指定 PPTX 可打开且严格为 25 页", 2, "hurdle"),
    "R002": ("源 PPT 删除指定三页后按原顺序完整插入第 9–23 页", 5, "hurdle"),
    "R003": ("源 PDF 的 10 页主题、正文和流程语义完整落到指定页", 5, "hurdle"),
    "R004": ("25 页文字、卡片、流程和图标均由原生可编辑对象组成", 5, "hurdle"),
    "R005": ("25 页均使用浅蓝背景且不存在深色整页遮罩", 3, "hurdle"),
    "R006": ("标题、正文、圆角框和视觉层级统一", 4, "positive"),
    "R007": ("右下角页码按最终顺序连续显示 1–25", 2, "positive"),
    "R008": ("LibreOffice 实际渲染为 25 页且关键文字可读", 2, "positive"),
    "R009": ("对象位于页面内且字号达到可读下限", 1, "positive"),
    "R010": ("出现整页截图、丢页、错误页序或不可编辑主体", -10, "penalty"),
    "R011": ("应删除的预约会议、工单列表或值守检查页仍残留", -5, "penalty"),
}

EXPECTED_TITLES = [
    "应急有线人工通信保障网",
    "目录",
    "网络定位与构成",
    "典型保障场景",
    "核心能力价值（一）",
    "核心能力价值（二）",
    "台站体系与协作架构",
    "典型设备组成",
    "云岚调度坐席平台",
    "内容结构",
    "系统总览：调度坐席的四层协作",
    "登录：进入值守环境",
    "主界面：把高频操作放在同一屏",
    "焦点规则：先确认“操作”，再确认“语音”",
    "工单窗口：先记录，再执行呼叫动作",
    "线路状态：用可视化面板降低误操作",
    "接续流程：从来话到三方沟通",
    "转接与交叉操作：语音不一定跟着窗口走",
    "合并与拆分：处理复杂通话关系",
    "查号：用更少输入找到正确联系人",
    "即时会议：临时协同但不污染会议列表",
    "座席互联：用于班组内快速沟通",
    "主菜单：按业务动作重新归类",
    "运行流程与管理要点",
    "谢谢聆听",
]

PDF_REQUIRED = {
    1: ["应急有线人工通信保障网", "定义、场景、价值与台站体系解析", "人工接续", "有线传输", "闭环记录"],
    2: ["目录", "网络定位与构成", "典型保障场景", "核心能力价值", "台站层级与设备"],
    3: ["什么是人工通信保障网", "人工接续台", "有线传输链路", "用户终端", "安全与记录", "人工判断", "身份核对"],
    4: ["应急指挥备份", "园区专线联络", "偏远站点接入", "跨班值守交接", "大型活动保障", "灾害抢险协同"],
    5: ["稳定性", "物理线路", "可控性", "权限判断", "责任闭环"],
    6: ["优先级调度", "快速恢复", "低依赖运行", "易培训维护"],
    7: ["中心汇接站", "区域转接站", "现场保障站", "终端接入点", "层级化管理"],
    8: ["人工交换坐席", "线路分配架", "有线传输介质", "安全记录设备", "值守终端"],
    24: ["接入确认", "记录需求", "选择通道", "呼出转接", "归档复核", "可靠接入", "身份核对", "闭环记录"],
    25: ["谢谢聆听", "稳定", "可控", "可恢复", "人工有线通信链路"],
}

SOURCE_KEEP_PAGES = list(range(1, 13)) + [14, 15, 17]
FORBIDDEN_TITLES = [
    "预约会议：把参会方和发起流程提前准备好",
    "工单列表与预约回呼：让未完成事项可追踪",
    "值守检查清单",
]


def criterion(identifier: str, passed: bool, evidence: str = "") -> dict:
    description, weight, kind = CRITERION_SPECS[identifier]
    return {
        "criterion_id": identifier,
        "description": description,
        "weight": weight,
        "type": kind,
        "passed": bool(passed),
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
    }


def normalize(text: str) -> str:
    return re.sub(r"\s+", "", text or "")


def slide_text(slide) -> str:
    return normalize(
        "\n".join(
            shape.text_frame.text
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False)
        )
    )


def text_counter_without_footer_page(slide, slide_width, slide_height) -> Counter:
    values = []
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        value = normalize(shape.text_frame.text)
        if not value:
            continue
        if (
            value.isdigit()
            and shape.left > slide_width * 0.85
            and shape.top > slide_height * 0.85
        ):
            continue
        values.append(value)
    return Counter(values)


def source_order_check(source: Presentation, candidate: Presentation) -> tuple[bool, str]:
    if len(source.slides) != 18 or len(candidate.slides) != 25:
        return False, f"页数异常：source={len(source.slides)} candidate={len(candidate.slides)}"
    errors = []
    for target_page, source_page in enumerate(SOURCE_KEEP_PAGES, start=9):
        expected = text_counter_without_footer_page(
            source.slides[source_page - 1], source.slide_width, source.slide_height
        )
        actual = text_counter_without_footer_page(
            candidate.slides[target_page - 1], candidate.slide_width, candidate.slide_height
        )
        if expected != actual:
            errors.append(
                f"p{target_page}≠源PPT p{source_page}：缺失={list((expected-actual).elements())[:4]}，新增={list((actual-expected).elements())[:4]}"
            )
    title_errors = []
    for page, title in enumerate(EXPECTED_TITLES, start=1):
        if normalize(title) not in slide_text(candidate.slides[page - 1]):
            title_errors.append(f"p{page}:{title}")
    if title_errors:
        errors.append(f"标题/页序不符={title_errors}")
    return not errors, "15 个源 PPT 页面文字逐页保持且标题顺序正确" if not errors else "；".join(errors)


def pdf_content_check(candidate: Presentation) -> tuple[bool, str]:
    missing = {}
    for page, phrases in PDF_REQUIRED.items():
        content = slide_text(candidate.slides[page - 1])
        absent = [phrase for phrase in phrases if normalize(phrase) not in content]
        if absent:
            missing[page] = absent
    return not missing, "PDF 十页的主题、正文、设备和流程关键语义完整" if not missing else f"缺失={missing}"


def editability_check(candidate: Presentation) -> tuple[bool, str, bool]:
    errors = []
    full_page_picture = False
    vector_types = {
        MSO_SHAPE_TYPE.AUTO_SHAPE,
        MSO_SHAPE_TYPE.FREEFORM,
        MSO_SHAPE_TYPE.GROUP,
        MSO_SHAPE_TYPE.LINE,
        MSO_SHAPE_TYPE.TEXT_BOX,
    }
    page_area = candidate.slide_width * candidate.slide_height
    for page, slide in enumerate(candidate.slides, start=1):
        pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        for picture in pictures:
            if picture.width * picture.height >= page_area * 0.75:
                full_page_picture = True
        if pictures:
            errors.append(f"p{page} 含 {len(pictures)} 个不可编辑栅格图片")
        editable_text = sum(
            1 for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and normalize(shape.text_frame.text)
        )
        editable_vector = sum(1 for shape in slide.shapes if shape.shape_type in vector_types)
        if editable_text < 3 or editable_vector < 6:
            errors.append(f"p{page} 可编辑文字/矢量对象不足：{editable_text}/{editable_vector}")
    return not errors, "25 页无栅格截图，文字与流程均为原生可编辑对象" if not errors else "；".join(errors), full_page_picture


def background_check(candidate: Presentation) -> tuple[bool, str]:
    errors = []
    page_area = candidate.slide_width * candidate.slide_height
    backgrounds = []
    for page, slide in enumerate(candidate.slides, start=1):
        value = None
        try:
            if slide.background.fill.type == 1 and slide.background.fill.fore_color.rgb:
                value = str(slide.background.fill.fore_color.rgb).upper()
        except Exception:
            value = None
        backgrounds.append(value)
        if not value or not upstream_verifier.is_light_blue(value):
            errors.append(f"p{page} 背景不是浅蓝：{value}")
        for shape in slide.shapes:
            if shape.width * shape.height < page_area * 0.90:
                continue
            try:
                fill = str(shape.fill.fore_color.rgb).upper() if shape.fill.type == 1 and shape.fill.fore_color.rgb else None
            except Exception:
                fill = None
            if fill and not upstream_verifier.is_light_blue(fill):
                errors.append(f"p{page} 存在深色/异色整页遮罩：{shape.name}={fill}")
    return not errors, f"背景={Counter(backgrounds)}" if not errors else "；".join(errors)


def style_check(candidate: Presentation) -> tuple[bool, str]:
    data = upstream_verifier.load_slide_data(candidate)
    passed, details = upstream_verifier.check_plus3_style(candidate, data)
    return bool(passed), "；".join(str(item) for item in details)


def page_number_check(candidate: Presentation) -> tuple[bool, str]:
    errors = []
    for page, slide in enumerate(candidate.slides, start=1):
        numbers = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            if shape.left <= candidate.slide_width * 0.85 or shape.top <= candidate.slide_height * 0.85:
                continue
            value = normalize(shape.text_frame.text)
            if value.isdigit():
                numbers.append(int(value))
        if numbers != [page]:
            errors.append(f"p{page}:{numbers}")
    return not errors, "右下角页码连续 1–25" if not errors else f"异常={errors}"


def render_check(path: Path) -> tuple[bool, str]:
    if not all(shutil.which(name) for name in ("libreoffice", "pdfinfo", "pdftotext")):
        return False, "缺少 libreoffice、pdfinfo 或 pdftotext"
    with tempfile.TemporaryDirectory(prefix="cowork-evo-068-") as directory:
        converted = subprocess.run(
            ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", directory, str(path)],
            capture_output=True,
            text=True,
            timeout=180,
            check=False,
        )
        pdf = Path(directory) / f"{path.stem}.pdf"
        if converted.returncode != 0 or not pdf.is_file():
            return False, (converted.stderr or converted.stdout or "PDF 转换失败")[-1000:]
        info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True, timeout=60, check=False)
        extracted = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], capture_output=True, text=True, timeout=60, check=False)
        match = re.search(r"^Pages:\s+(\d+)", info.stdout, flags=re.M)
        pages = int(match.group(1)) if match else 0
        content = normalize(extracted.stdout)
    required = ["应急有线人工通信保障网", "典型设备组成", "云岚调度坐席平台", "即时会议", "主菜单", "运行流程与管理要点", "谢谢聆听"]
    missing = [value for value in required if normalize(value) not in content]
    return pages == 25 and not missing, f"渲染页数={pages}，缺失={missing}"


def professional_check(candidate: Presentation) -> tuple[bool, str]:
    errors = []
    tolerance = 50000
    small_fonts = []
    for page, slide in enumerate(candidate.slides, start=1):
        for shape in slide.shapes:
            if shape.left < -tolerance or shape.top < -tolerance or shape.left + shape.width > candidate.slide_width + tolerance or shape.top + shape.height > candidate.slide_height + tolerance:
                errors.append(f"p{page}:{shape.name} 超出页面")
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip() and run.font.size is not None and run.font.size.pt < 8:
                        small_fonts.append(f"p{page}:{run.font.size.pt}")
    if small_fonts:
        errors.append(f"小于 8pt={small_fonts[:10]}")
    return not errors, "所有对象位于页内，显式字号不低于 8pt" if not errors else "；".join(errors)


def evaluate(artifact_dir: Path) -> dict:
    output = artifact_dir / OUTPUT_NAME
    results = []
    if not output.is_file():
        results.append(criterion("R001", False, f"缺少 {OUTPUT_NAME}"))
        return build_result(results, output)
    try:
        candidate = Presentation(output)
    except Exception as exc:
        results.append(criterion("R001", False, f"文件无法打开：{exc}"))
        return build_result(results, output)

    valid = len(candidate.slides) == 25
    results.append(criterion("R001", valid, f"页数={len(candidate.slides)}"))
    if not valid:
        results.append(criterion("R010", True, "页数损坏或合并未完成"))
        return build_result(results, output)
    task_root = Path(__file__).resolve().parents[1]
    source = Presentation(task_root / "data" / "input_files" / SOURCE_PPT)

    ordered, order_evidence = source_order_check(source, candidate)
    results.append(criterion("R002", ordered, order_evidence))
    pdf_ok, pdf_evidence = pdf_content_check(candidate)
    results.append(criterion("R003", pdf_ok, pdf_evidence))
    editable, editable_evidence, full_page_picture = editability_check(candidate)
    results.append(criterion("R004", editable, editable_evidence))
    background_ok, background_evidence = background_check(candidate)
    results.append(criterion("R005", background_ok, background_evidence))
    styled, style_evidence = style_check(candidate)
    results.append(criterion("R006", styled, style_evidence))
    numbered, page_evidence = page_number_check(candidate)
    results.append(criterion("R007", numbered, page_evidence))
    rendered, render_evidence = render_check(output)
    results.append(criterion("R008", rendered, render_evidence))
    professional, professional_evidence = professional_check(candidate)
    results.append(criterion("R009", professional, professional_evidence))

    structural_hit = not valid or not ordered or not editable or full_page_picture
    results.append(criterion("R010", structural_hit, "主体截图化、丢页或顺序损坏" if structural_hit else "页面顺序与可编辑结构正常"))
    all_text = normalize("\n".join(slide_text(slide) for slide in candidate.slides))
    leaked = [title for title in FORBIDDEN_TITLES if normalize(title) in all_text]
    results.append(criterion("R011", bool(leaked), f"残留={leaked}"))
    return build_result(results, output)


def build_result(results: list[dict], artifact_path: Path) -> dict:
    present = {item["criterion_id"] for item in results}
    for identifier in CRITERION_SPECS:
        if identifier not in present:
            results.append(criterion(identifier, False, "因交付不存在或损坏，本判据未执行"))
    results.sort(key=lambda item: item["criterion_id"])
    positive_total = sum(item["weight"] for item in results if item["weight"] > 0)
    earned = sum(item["weight"] for item in results if item["weight"] > 0 and item["passed"])
    penalties = sum(item["weight"] for item in results if item["weight"] < 0 and item["passed"])
    reward = max(0.0, min(1.0, (earned + penalties) / positive_total))
    hurdles = [item["criterion_id"] for item in results if item["type"] == "hurdle" and not item["passed"]]
    if hurdles:
        reward = 0.0
    digest = hashlib.sha256(artifact_path.read_bytes()).hexdigest() if artifact_path.is_file() else hashlib.sha256(b"").hexdigest()
    normalized_results = []
    for item in results:
        score = (0.0 if item["passed"] else 1.0) if item["type"] == "penalty" else item["score"]
        normalized_results.append({
            "criterion_id": item["criterion_id"],
            "score": score,
            "evidence": item["evidence"],
            "method": item["method"],
            "raw": {"description": item["description"], "weight": item["weight"], "type": item["type"], "passed": item["passed"]},
        })
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [digest],
        "criterion_results": normalized_results,
        "reward": round(reward, 6),
        "passed": reward >= PASS_THRESHOLD,
        "hurdle_failures": hurdles,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    args = parser.parse_args()
    result = evaluate(args.artifact_dir)
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
