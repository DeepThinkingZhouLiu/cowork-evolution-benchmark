#!/usr/bin/env python3
"""officeval_069 的逐题参考解法。

这不是通用 PPT 生成器：它只处理这道真实输入中由参考截图标注的
10 个圆角矩形，其他 54 个对象均不改动。
"""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Pt


INPUT_NAME = "工商层面_流程图.pptx"
OUTPUT_NAME = "工商层面_流程图_框体字体调整版.pptx"


TARGETS = [
    {
        "keys": ("2023年", "管理负责人调整", "职能重新划分"),
        "text": "2023年\n管理负责人调整\n职能重新划分",
        "width_cm": 5.08,
        "height_cm": 2.11,
        "fill": "DDE9FF",
        "line": "3768B3",
        "font_pt": 13.2,
    },
    {
        "keys": ("2025年", "引入合规与数据", "治理机制"),
        "text": "2025年\n引入合规与数据\n治理机制",
        "width_cm": 5.08,
        "height_cm": 2.11,
        "fill": "DDE9FF",
        "line": "3768B3",
        "font_pt": 13.2,
    },
    {
        "keys": ("林远航", "持股 65%", "发起投资"),
        "text": "林远航\n持股 65%\n发起投资",
        "width_cm": 4.45,
        "height_cm": 2.08,
        "fill": "DDE9FF",
        "line": "3768B3",
        "font_pt": 12.3,
    },
    {
        "keys": ("周予安", "持股 35%", "产业资源"),
        "text": "周予安\n持股 35%\n产业资源",
        "width_cm": 4.45,
        "height_cm": 2.08,
        "fill": "DDE9FF",
        "line": "3768B3",
        "font_pt": 12.3,
    },
    {
        "keys": ("周予安", "持股 30%", "总经理"),
        "text": "周予安\n持股 30%\n总经理",
        "width_cm": 4.45,
        "height_cm": 2.08,
        "fill": "DDE9FF",
        "line": "3768B3",
        "font_pt": 12.3,
    },
    {
        "keys": ("董事会决策层", "战略审批", "重大合同"),
        "text": "董事会决策层\n战略审批 / 重大合同",
        "width_cm": 4.45,
        "height_cm": 2.08,
        "fill": "DDE9FF",
        "line": "3768B3",
        "font_pt": 12.3,
    },
    {
        "keys": ("仓配试点小组", "接入订单", "库位", "车辆数据"),
        "text": "仓配试点小组\n接入订单、库位、\n车辆数据",
        "width_cm": 4.45,
        "height_cm": 2.08,
        "fill": "F8E3E8",
        "line": "AD4E66",
        "font_pt": 12.0,
    },
    {
        "keys": ("日常运营链路", "订单调度", "仓储履约", "异常复盘"),
        "text": "日常运营链路\n订单调度 → 仓储履约\n→ 异常复盘",
        "width_cm": 4.45,
        "height_cm": 2.08,
        "fill": "F8E3E8",
        "line": "AD4E66",
        "font_pt": 12.0,
    },
    {
        "keys": ("经营执行线", "周予安", "总经理"),
        "text": "经营执行线\n周予安\n总经理",
        "width_cm": 4.45,
        "height_cm": 2.08,
        "fill": "F8E3E8",
        "line": "AD4E66",
        "font_pt": 12.0,
    },
    {
        "keys": ("监督复核线", "许知澜", "监事"),
        "text": "监督复核线\n许知澜\n监事",
        "width_cm": 4.45,
        "height_cm": 2.08,
        "fill": "F8E3E8",
        "line": "AD4E66",
        "font_pt": 12.0,
    },
]


def find_target(slide, keys: tuple[str, ...]):
    matches = []
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        text = shape.text_frame.text
        if all(key in text for key in keys):
            matches.append(shape)
    if len(matches) != 1:
        raise ValueError(f"目标 {keys} 应唯一匹配，实际匹配 {len(matches)} 个")
    return matches[0]


def set_box_style(shape, spec: dict) -> None:
    center_x = shape.left + shape.width / 2
    center_y = shape.top + shape.height / 2
    width = Cm(spec["width_cm"])
    height = Cm(spec["height_cm"])
    shape.left = int(center_x - width / 2)
    shape.top = int(center_y - height / 2)
    shape.width = width
    shape.height = height

    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(spec["fill"])
    shape.line.color.rgb = RGBColor.from_string(spec["line"])
    shape.line.width = Pt(1.2)

    text_frame = shape.text_frame
    text_frame.text = spec["text"]
    text_frame.word_wrap = True
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    text_frame.margin_left = Cm(0.12)
    text_frame.margin_right = Cm(0.12)
    text_frame.margin_top = Cm(0.05)
    text_frame.margin_bottom = Cm(0.05)
    for paragraph in text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER
        for run in paragraph.runs:
            run.font.name = "Noto Sans CJK SC"
            run.font.size = Pt(spec["font_pt"])
            run.font.bold = True


def solve(input_path: Path, output_path: Path) -> None:
    presentation = Presentation(input_path)
    if len(presentation.slides) != 1:
        raise ValueError(f"预期 1 页，实际 {len(presentation.slides)} 页")
    slide = presentation.slides[0]
    if len(slide.shapes) != 64:
        raise ValueError(f"预期 64 个对象，实际 {len(slide.shapes)} 个")
    for spec in TARGETS:
        set_box_style(find_target(slide, spec["keys"]), spec)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=Path("/app/data/input_files"))
    parser.add_argument("--output-dir", type=Path, default=Path("/app/output"))
    args = parser.parse_args()
    solve(args.input_dir / INPUT_NAME, args.output_dir / OUTPUT_NAME)


if __name__ == "__main__":
    main()
