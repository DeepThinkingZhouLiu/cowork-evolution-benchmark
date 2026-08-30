#!/usr/bin/env python3
"""officeval_069 的 Harbor Judge。

在保留上游逐题 Verifier 的基础上，增加了未标注对象保护、
可编辑性、防截图覆盖和 LibreOffice 实际渲染检查。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import upstream_verifier


OUTPUT_NAME = "工商层面_流程图_框体字体调整版.pptx"
INPUT_NAME = "工商层面_流程图.pptx"
TASK_ID = "cowork-evo/officeval-069-reference-guided-visual-editing"
RUBRIC_VERSION = "omegause-officeval-069-adaptation-v0.1"
JUDGE_VERSION = "officeval-069-adaptation-v0.1"
EMU_PER_CM = 360000.0

CRITERION_SPECS = {
    "R001": ("指定 PPTX 可打开，且页数与页面尺寸不变", 2, "hurdle"),
    "R002": ("64 个流程图对象和 10 个目标框保持可编辑", 2, "hurdle"),
    "R003": ("未标注的 54 个对象及目标文字内容保持不变", 3, "positive"),
    "R004": ("2023 年时间轴框样式正确", 1, "positive"),
    "R005": ("2025 年时间轴框样式正确", 1, "positive"),
    "R006": ("林远航 65% 发起投资框样式正确", 1, "positive"),
    "R007": ("周予安 35% 产业资源框样式正确", 1, "positive"),
    "R008": ("周予安 30% 总经理框样式正确", 1, "positive"),
    "R009": ("董事会决策层框样式正确", 1, "positive"),
    "R010": ("仓配试点小组框样式正确", 1, "positive"),
    "R011": ("日常运营链路框样式正确", 1, "positive"),
    "R012": ("经营执行线框样式正确", 1, "positive"),
    "R013": ("监督复核线框样式正确", 1, "positive"),
    "R014": ("LibreOffice 实际渲染为单页且关键文字可读", 2, "positive"),
    "R015": ("使用大面积截图覆盖或把流程图压平", -8, "penalty"),
    "R016": ("两个及以上未标注对象被删除、改文或移动", -6, "penalty"),
}


def criterion(identifier: str, passed: bool, evidence: str = "") -> dict:
    description, weight, kind = CRITERION_SPECS[identifier]
    return {
        "criterion_id": identifier,
        "description": description,
        "weight": weight,
        "type": kind,
        "passed": bool(passed),
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
    }


def normalized_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return re.sub(r"\s+", "", shape.text_frame.text)


def explicit_color(color_format) -> str | None:
    try:
        return str(color_format.rgb)
    except Exception:
        return None


def fill_color(shape) -> str | None:
    try:
        return explicit_color(shape.fill.fore_color)
    except Exception:
        return None


def line_color(shape) -> str | None:
    try:
        return explicit_color(shape.line.color)
    except Exception:
        return None


def font_signature(shape) -> tuple:
    if not getattr(shape, "has_text_frame", False):
        return ()
    values = []
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            values.append(
                (
                    run.font.name,
                    round(run.font.size.pt, 2) if run.font.size is not None else None,
                    run.font.bold,
                )
            )
    return tuple(values)


def preset_geometry(shape) -> str | None:
    """读取 OOXML prst 名，避免 python-pptx 遇到 prst=line 时抛错。"""
    try:
        preset = shape._element.spPr.prstGeom
        return preset.get("prst") if preset is not None else None
    except Exception:
        return None


def find_rule_indexes(slide) -> tuple[dict[int, int], list[str]]:
    indexes = {}
    errors = []
    for rule_index, rule in enumerate(upstream_verifier.PLUS_RULES):
        matches = []
        for shape_index, shape in enumerate(slide.shapes):
            if not getattr(shape, "has_text_frame", False):
                continue
            if all(key in shape.text_frame.text for key in rule["keys"]):
                matches.append(shape_index)
        if len(matches) != 1:
            errors.append(f"规则{rule_index + 1}匹配 {len(matches)} 个对象")
        else:
            indexes[rule_index] = matches[0]
    return indexes, errors


def compare_preservation(source: Presentation, candidate: Presentation) -> list[str]:
    errors = []
    if len(source.slides) != 1 or len(candidate.slides) != 1:
        return ["页数不是 1"]
    before_slide, after_slide = source.slides[0], candidate.slides[0]
    if len(before_slide.shapes) != len(after_slide.shapes):
        errors.append(f"对象数 {len(before_slide.shapes)} -> {len(after_slide.shapes)}")

    target_indexes, target_lookup_errors = find_rule_indexes(before_slide)
    errors.extend(target_lookup_errors)
    target_shape_indexes = set(target_indexes.values())
    for index, (before, after) in enumerate(zip(before_slide.shapes, after_slide.shapes)):
        if before.shape_type != after.shape_type:
            errors.append(f"对象{index}类型改变")
            continue
        if normalized_text(before) != normalized_text(after):
            label = "目标" if index in target_shape_indexes else "非目标"
            errors.append(f"{label}对象{index}文字改变")
        if index in target_shape_indexes:
            continue

        geometry_delta = max(
            abs(getattr(before, field) - getattr(after, field)) / EMU_PER_CM
            for field in ("left", "top", "width", "height")
        )
        if geometry_delta > 0.03:
            errors.append(f"非目标对象{index}位置或尺寸改变 {geometry_delta:.3f}cm")
        if preset_geometry(before) != preset_geometry(after):
            errors.append(f"非目标对象{index}形状改变")
        if fill_color(before) != fill_color(after):
            errors.append(f"非目标对象{index}填充颜色改变")
        if line_color(before) != line_color(after):
            errors.append(f"非目标对象{index}框线颜色改变")
        if font_signature(before) != font_signature(after):
            errors.append(f"非目标对象{index}字体样式改变")
    return sorted(set(errors))


def editability_check(candidate: Presentation) -> tuple[bool, list[str], bool]:
    errors = []
    slide = candidate.slides[0]
    rule_indexes, rule_errors = find_rule_indexes(slide)
    errors.extend(rule_errors)
    if len(slide.shapes) != 64:
        errors.append(f"对象数应为 64，实际 {len(slide.shapes)}")
    for rule_index, shape_index in rule_indexes.items():
        shape = slide.shapes[shape_index]
        if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE or not shape.has_text_frame:
            errors.append(f"规则{rule_index + 1}的目标不是可编辑文字形状")

    page_area = candidate.slide_width * candidate.slide_height
    large_pictures = []
    for index, shape in enumerate(slide.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            ratio = (shape.width * shape.height) / page_area
            if ratio >= 0.60:
                large_pictures.append(f"图片对象{index}覆盖{ratio:.0%}页面")
    errors.extend(large_pictures)
    flattened = bool(large_pictures) or len(rule_indexes) < 10 or len(slide.shapes) < 54
    return not errors, errors, flattened


def render_check(path: Path) -> tuple[bool, str]:
    libreoffice = shutil.which("libreoffice")
    pdftotext = shutil.which("pdftotext")
    if not libreoffice or not pdftotext:
        return False, "缺少 libreoffice 或 pdftotext"
    with tempfile.TemporaryDirectory(prefix="cowork-evo-069-") as temp_dir:
        converted = subprocess.run(
            [libreoffice, "--headless", "--convert-to", "pdf", "--outdir", temp_dir, str(path)],
            capture_output=True,
            text=True,
            timeout=180,
            check=False,
        )
        pdf_path = Path(temp_dir) / f"{path.stem}.pdf"
        if converted.returncode != 0 or not pdf_path.is_file():
            return False, (converted.stderr or converted.stdout or "PDF 转换失败")[-800:]
        extracted = subprocess.run(
            [pdftotext, "-layout", str(pdf_path), "-"],
            capture_output=True,
            text=True,
            timeout=60,
            check=False,
        )
        if extracted.returncode != 0:
            return False, (extracted.stderr or "PDF 文字抽取失败")[-800:]
        pages = [page for page in extracted.stdout.split("\f") if page.strip()]
        normalized = re.sub(r"\s+", "", "".join(pages))
    # pdftotext 会按整页横向阅读顺序交错多列文本，因此逐个核对
    # 上游规则中的关键片段，不错误要求同一框的文本在 PDF 流中相邻。
    required = sorted(
        {
            re.sub(r"\s+", "", key)
            for rule in upstream_verifier.PLUS_RULES
            for key in rule["keys"]
        }
    )
    missing = [text for text in required if text not in normalized]
    return len(pages) == 1 and not missing, f"页数={len(pages)}，缺失文字={missing}"


def evaluate(artifact_dir: Path) -> dict:
    output_path = artifact_dir / OUTPUT_NAME
    results = []
    if not output_path.is_file():
        results.append(criterion("R001", False, f"缺少 {OUTPUT_NAME}"))
        return build_result(results, output_path)
    try:
        candidate = Presentation(output_path)
    except Exception as exc:
        results.append(criterion("R001", False, f"文件无法打开：{exc}"))
        return build_result(results, output_path)

    width_cm = candidate.slide_width / EMU_PER_CM
    height_cm = candidate.slide_height / EMU_PER_CM
    file_ok = (
        len(candidate.slides) == 1
        and abs(width_cm - 50.80) <= 0.05
        and abs(height_cm - 28.575) <= 0.05
    )
    results.append(
        criterion(
            "R001",
            file_ok,
            f"页数={len(candidate.slides)}，尺寸={width_cm:.3f}×{height_cm:.3f}cm",
        )
    )
    if not file_ok:
        return build_result(results, output_path)

    editable, edit_errors, flattened = editability_check(candidate)
    results.append(
        criterion(
            "R002",
            editable,
            "64 个对象及 10 个目标均可编辑" if editable else "；".join(edit_errors[:20]),
        )
    )

    source_path = Path(os.environ.get("INPUT_DIR", "/app/input_files")) / INPUT_NAME
    preservation_errors = []
    try:
        source = Presentation(source_path)
        preservation_errors = compare_preservation(source, candidate)
    except Exception as exc:
        preservation_errors = [f"无法比较原始输入：{exc}"]
    results.append(
        criterion(
            "R003",
            not preservation_errors,
            "未标注对象和文字内容完整保留"
            if not preservation_errors
            else "；".join(preservation_errors[:30]),
        )
    )

    slide = candidate.slides[0]
    for criterion_id, rule in zip(
        ("R004", "R005", "R006", "R007", "R008", "R009", "R010", "R011", "R012", "R013"),
        upstream_verifier.PLUS_RULES,
    ):
        passed, reasons = upstream_verifier.check_plus_rule(slide, rule)
        evidence = "上游逐题规则通过" if passed else "；".join(reasons)
        results.append(criterion(criterion_id, passed, evidence))

    rendered, render_evidence = render_check(output_path)
    results.append(criterion("R014", rendered, render_evidence))
    results.append(criterion("R015", flattened, "；".join(edit_errors) if flattened else "未发现截图覆盖或压平"))
    results.append(
        criterion(
            "R016",
            len(preservation_errors) >= 2,
            f"保留性错误数={len(preservation_errors)}",
        )
    )
    return build_result(results, output_path)


def build_result(results: list[dict], artifact_path: Path) -> dict:
    present = {item["criterion_id"] for item in results}
    for identifier in CRITERION_SPECS:
        if identifier not in present:
            results.append(criterion(identifier, False, "因交付不存在、损坏或门槛失败，本判据未执行"))
    results.sort(key=lambda item: item["criterion_id"])

    positive_total = sum(item["weight"] for item in results if item["weight"] > 0)
    earned = sum(item["weight"] for item in results if item["weight"] > 0 and item["passed"])
    penalties = sum(item["weight"] for item in results if item["weight"] < 0 and item["passed"])
    reward = max(0.0, min(1.0, (earned + penalties) / positive_total))
    hurdle_failures = [
        item["criterion_id"] for item in results if item["type"] == "hurdle" and not item["passed"]
    ]
    if hurdle_failures:
        reward = 0.0

    digest = hashlib.sha256(artifact_path.read_bytes()).hexdigest() if artifact_path.is_file() else hashlib.sha256(b"").hexdigest()
    normalized = []
    for item in results:
        score = (0.0 if item["passed"] else 1.0) if item["type"] == "penalty" else item["score"]
        normalized.append(
            {
                "criterion_id": item["criterion_id"],
                "score": score,
                "evidence": item["evidence"],
                "method": item["method"],
                "raw": {
                    "description": item["description"],
                    "weight": item["weight"],
                    "type": item["type"],
                    "passed": item["passed"],
                },
            }
        )
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [digest],
        "criterion_results": normalized,
        "reward": round(reward, 6),
        "passed": reward >= 0.75,
        "hurdle_failures": hurdle_failures,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path)
    args = parser.parse_args()
    result = evaluate(args.artifact_dir)
    rendered = json.dumps(result, ensure_ascii=False, indent=2)
    print(rendered)
    if args.result:
        args.result.parent.mkdir(parents=True, exist_ok=True)
        args.result.write_text(rendered + "\n", encoding="utf-8")
    raise SystemExit(0 if result["passed"] else 1)


if __name__ == "__main__":
    main()
