#!/usr/bin/env python3
"""officeval_070 的逐题参考解法。

只调整原始单页 PPTX 中六个研究板块的编号、标题和正文；不生成新内容，
也不作为其他题目的通用制题或版式生成器。
"""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Cm, Pt


INPUT_NAME = "技术路线图.pptx"
OUTPUT_NAME = "技术路线图_调整美化版.pptx"

SECTIONS = [
    ("01", "研究内容一：关键病理事件界定", "临床分层与样本库建立"),
    ("02", "研究内容二：色氨酸代谢异常定位", "卵泡液与颗粒细胞代谢组检测"),
    ("03", "研究内容三：ATF4核内调控解析", "ATF4表达、磷酸化与核定位检测"),
    ("04", "研究内容四：表观遗传持续化机制", "H3K4me3与H3K27ac开放状态检测"),
    ("05", "研究内容五：轴向因果验证", "IA梯度干预与时间窗干预"),
    ("06", "预期结果与科学意义", "揭示PCOS持续性排卵障碍的新病理机制"),
]


def shape_text(shape) -> str:
    return shape.text.strip() if getattr(shape, "has_text_frame", False) else ""


def find_unique_text_shape(slide, predicate, label: str):
    matches = [shape for shape in slide.shapes if predicate(shape_text(shape))]
    if len(matches) != 1:
        raise ValueError(f"{label} 应唯一匹配，实际找到 {len(matches)} 个")
    return matches[0]


def collect_frames(slide) -> list:
    """按行、列收集六个研究板块的底层流程框图片。"""

    frames = []
    for shape in slide.shapes:
        if shape.shape_type != 13:
            continue
        width_cm = shape.width / 360000
        height_cm = shape.height / 360000
        top_cm = shape.top / 360000
        if 7.5 <= width_cm <= 7.8 and 2.6 <= height_cm <= 2.9 and top_cm >= 12.0:
            frames.append(shape)
    frames.sort(key=lambda shape: (round(shape.top / 360000, 1), shape.left))
    if len(frames) != 6:
        raise ValueError(f"应找到 6 个研究板块底框，实际找到 {len(frames)} 个")
    # 几何顺序是左1、右1、左2、右2、左3、右3；转成 01、02、03、04、05、06。
    return [frames[0], frames[2], frames[4], frames[1], frames[3], frames[5]]


def set_geometry(shape, *, left, top, width, height) -> None:
    shape.left = int(left)
    shape.top = int(top)
    shape.width = int(width)
    shape.height = int(height)


def format_number(shape, frame) -> None:
    set_geometry(
        shape,
        left=frame.left + Cm(0.28),
        top=frame.top + Cm(0.12),
        width=Cm(0.70),
        height=Cm(0.65),
    )
    text_frame = shape.text_frame
    text_frame.word_wrap = True
    text_frame.auto_size = MSO_AUTO_SIZE.NONE
    text_frame.vertical_anchor = MSO_ANCHOR.TOP
    text_frame.margin_left = text_frame.margin_right = 0
    text_frame.margin_top = text_frame.margin_bottom = 0
    for paragraph in text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.line_spacing = 1.0
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "微软雅黑"
            run.font.size = Pt(11.0)
            run.font.bold = True


def format_title(shape, frame) -> None:
    set_geometry(
        shape,
        left=frame.left + Cm(1.05),
        top=frame.top + Cm(0.12),
        width=frame.width - Cm(1.40),
        height=Cm(0.65),
    )
    text_frame = shape.text_frame
    text_frame.word_wrap = True
    text_frame.auto_size = MSO_AUTO_SIZE.NONE
    text_frame.vertical_anchor = MSO_ANCHOR.TOP
    text_frame.margin_left = text_frame.margin_right = 0
    text_frame.margin_top = text_frame.margin_bottom = 0
    for paragraph in text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.line_spacing = 1.0
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "微软雅黑"
            run.font.size = Pt(11.0)
            run.font.bold = True


def format_body(shape, frame) -> None:
    set_geometry(
        shape,
        left=frame.left + Cm(0.35),
        top=frame.top + Cm(1.03),
        width=frame.width - Cm(0.70),
        height=frame.height - Cm(1.25),
    )
    text_frame = shape.text_frame
    text_frame.word_wrap = True
    text_frame.auto_size = MSO_AUTO_SIZE.NONE
    text_frame.vertical_anchor = MSO_ANCHOR.TOP
    text_frame.margin_left = text_frame.margin_right = 0
    text_frame.margin_top = text_frame.margin_bottom = 0
    for paragraph in text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.line_spacing = 0.95
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "微软雅黑"
            run.font.size = Pt(8.0)
            run.font.bold = False


def solve(input_dir: Path, output_dir: Path) -> Path:
    input_path = input_dir / INPUT_NAME
    if not input_path.is_file():
        raise FileNotFoundError(f"缺少输入文件：{input_path}")

    presentation = Presentation(input_path)
    if len(presentation.slides) != 1:
        raise ValueError(f"输入应为 1 页，实际为 {len(presentation.slides)} 页")
    slide = presentation.slides[0]
    frames = collect_frames(slide)

    for (number, title, body_prefix), frame in zip(SECTIONS, frames):
        number_shape = find_unique_text_shape(slide, lambda text, n=number: text == n, f"编号 {number}")
        title_shape = find_unique_text_shape(slide, lambda text, t=title: text == t, f"标题 {title}")
        body_shape = find_unique_text_shape(slide, lambda text, p=body_prefix: text.startswith(p), f"正文 {body_prefix}")
        format_number(number_shape, frame)
        format_title(title_shape, frame)
        format_body(body_shape, frame)

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / OUTPUT_NAME
    presentation.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    print(solve(args.input_dir, args.output_dir))


if __name__ == "__main__":
    main()
