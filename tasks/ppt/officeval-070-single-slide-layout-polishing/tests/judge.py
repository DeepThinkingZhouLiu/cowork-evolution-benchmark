#!/usr/bin/env python3
"""officeval_070 的 Harbor Judge。

保留并复用上游逐题 Verifier，同时补齐页面尺寸、非目标对象保护、
几何重叠、LibreOffice 实际渲染和本仓库统一结果 schema。
"""

from __future__ import annotations

import argparse
import hashlib
import itertools
import json
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation

import upstream_verifier


OUTPUT_NAME = "技术路线图_调整美化版.pptx"
INPUT_NAME = "技术路线图.pptx"
TASK_ID = "cowork-evo/officeval-070-single-slide-layout-polishing"
RUBRIC_VERSION = "omegause-officeval-070-adaptation-v0.1"
JUDGE_VERSION = "officeval-070-adaptation-v0.1"
EMU_PER_CM = 360000.0

SECTIONS = [
    ("01", "研究内容一：关键病理事件界定", "临床分层与样本库建立", "+1#研究内容一"),
    ("02", "研究内容二：色氨酸代谢异常定位", "卵泡液与颗粒细胞代谢组检测", "+1#研究内容二"),
    ("03", "研究内容三：ATF4核内调控解析", "ATF4表达、磷酸化与核定位检测", "+1#研究内容三"),
    ("04", "研究内容四：表观遗传持续化机制", "H3K4me3与H3K27ac开放状态检测", "+1#研究内容四"),
    ("05", "研究内容五：轴向因果验证", "IA梯度干预与时间窗干预", "+1#研究内容五"),
    ("06", "预期结果与科学意义", "揭示PCOS持续性排卵障碍的新病理机制", "+1#预期结果"),
]

CRITERION_SPECS = {
    "R001": ("交付文件可打开，并保持 1 页和原页面尺寸", 2, "hurdle"),
    "R002": ("研究内容一完整位于对应底框", 1, "positive"),
    "R003": ("研究内容二完整位于对应底框", 1, "positive"),
    "R004": ("研究内容三完整位于对应底框", 1, "positive"),
    "R005": ("研究内容四完整位于对应底框", 1, "positive"),
    "R006": ("研究内容五完整位于对应底框", 1, "positive"),
    "R007": ("预期结果完整位于对应底框", 1, "positive"),
    "R008": ("标题和正文字体满足统一规范", 3, "positive"),
    "R009": ("正文行距、边距和两列对齐满足统一规范", 3, "positive"),
    "R010": ("目标文字不互相覆盖，也不压住其他对象", 3, "positive"),
    "R011": ("非目标内容、对象和位置保持不变", 3, "positive"),
    "R012": ("LibreOffice 实际渲染为一页且内容可读", 2, "positive"),
    "R013": ("三个及以上研究板块仍有明显重叠或出框", -5, "penalty"),
    "R014": ("两个及以上非目标对象被破坏", -3, "penalty"),
}


def criterion(identifier: str, passed: bool, evidence: str = "") -> dict:
    description, weight, kind = CRITERION_SPECS[identifier]
    return {
        "criterion_id": identifier,
        "description": description,
        "weight": weight,
        "type": kind,
        "passed": bool(passed),
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
    }


def shape_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return "\n".join(paragraph.text for paragraph in shape.text_frame.paragraphs).strip()


def rect_cm(shape) -> tuple[float, float, float, float]:
    return tuple(value / EMU_PER_CM for value in (shape.left, shape.top, shape.width, shape.height))


def rect_contains(outer, inner, tolerance: float = 0.02) -> bool:
    ol, ot, ow, oh = outer
    il, it, iw, ih = inner
    return (
        il + tolerance >= ol
        and it + tolerance >= ot
        and il + iw - tolerance <= ol + ow
        and it + ih - tolerance <= ot + oh
    )


def overlap_area(first, second) -> float:
    fl, ft, fw, fh = first
    sl, st, sw, sh = second
    width = max(0.0, min(fl + fw, sl + sw) - max(fl, sl))
    height = max(0.0, min(ft + fh, st + sh) - max(ft, st))
    return width * height


def find_unique(slide, predicate, label: str):
    matches = [shape for shape in slide.shapes if predicate(shape_text(shape))]
    if len(matches) != 1:
        raise ValueError(f"{label} 应唯一匹配，实际 {len(matches)} 个")
    return matches[0]


def target_shape_indexes(slide) -> set[int]:
    indexes = set()
    for index, shape in enumerate(slide.shapes):
        text = shape_text(shape)
        if text in {item[0] for item in SECTIONS}:
            indexes.add(index)
        if text in {item[1] for item in SECTIONS}:
            indexes.add(index)
        if any(text.startswith(item[2]) for item in SECTIONS):
            indexes.add(index)
    return indexes


def compare_preserved_objects(source: Presentation, candidate: Presentation) -> list[str]:
    errors = []
    if len(source.slides) != len(candidate.slides):
        return [f"页数改变：{len(source.slides)} -> {len(candidate.slides)}"]
    source_slide, candidate_slide = source.slides[0], candidate.slides[0]
    if len(source_slide.shapes) != len(candidate_slide.shapes):
        errors.append(f"对象数改变：{len(source_slide.shapes)} -> {len(candidate_slide.shapes)}")
    target_indexes = target_shape_indexes(source_slide)
    for index, (before, after) in enumerate(zip(source_slide.shapes, candidate_slide.shapes)):
        if before.shape_type != after.shape_type:
            errors.append(f"对象{index}类型改变")
            continue
        if index in target_indexes:
            if shape_text(before) != shape_text(after):
                errors.append(f"目标对象{index}文字被改写")
            continue
        if shape_text(before) != shape_text(after):
            errors.append(f"非目标对象{index}文字改变")
        if (before.left, before.top, before.width, before.height) != (after.left, after.top, after.width, after.height):
            errors.append(f"非目标对象{index}位置或尺寸改变")
        if before.shape_type == 13:
            try:
                if before.image.sha1 != after.image.sha1:
                    errors.append(f"非目标图片对象{index}内容改变")
            except Exception as exc:
                errors.append(f"非目标图片对象{index}无法比较：{exc}")
    return errors


def geometry_check(slide) -> tuple[list[str], int]:
    method_top = upstream_verifier.find_method_title_top(slide)
    if method_top is None:
        return ["缺少“研究方法与技术路线”标题"], 6
    rows = upstream_verifier.collect_section_frames(slide, method_top)
    issues = []
    severe_sections = 0
    all_targets = []

    for number, title, body_prefix, _ in SECTIONS:
        title_shape = find_unique(slide, lambda text, expected=title: text == expected, title)
        body_shape = find_unique(slide, lambda text, prefix=body_prefix: text.startswith(prefix), body_prefix)
        number_shape = find_unique(slide, lambda text, expected=number: text == expected, number)
        all_targets.extend([number_shape, title_shape, body_shape])

        section_spec = next(item for item in upstream_verifier.SECTIONS if item[1] == title)
        frame = upstream_verifier.locate_frame(rows, section_spec[2], section_spec[3])
        section_issues = []
        if frame is None:
            section_issues.append("找不到对应底框")
        else:
            frame_rect = frame["rect"]
            for label, shape in (("编号", number_shape), ("标题", title_shape), ("正文", body_shape)):
                if not rect_contains(frame_rect, rect_cm(shape), tolerance=0.20):
                    section_issues.append(f"{label}超出底框")
        for (left_label, left_shape), (right_label, right_shape) in itertools.combinations(
            (("编号", number_shape), ("标题", title_shape), ("正文", body_shape)), 2
        ):
            area = overlap_area(rect_cm(left_shape), rect_cm(right_shape))
            if area > 0.01:
                section_issues.append(f"{left_label}与{right_label}重叠 {area:.3f}cm²")
        if section_issues:
            severe_sections += 1
            issues.append(f"{number}: " + "；".join(section_issues))

    # 六组文本之间也必须互不覆盖。
    for first, second in itertools.combinations(all_targets, 2):
        area = overlap_area(rect_cm(first), rect_cm(second))
        if area > 0.01 and shape_text(first)[:2] != shape_text(second)[:2]:
            issues.append(f"跨板块重叠：{shape_text(first)[:12]} / {shape_text(second)[:12]}，{area:.3f}cm²")

    # 排除整页背景和承载目标文本的底框，检查是否压住箭头等邻接图片。
    for target in all_targets:
        target_rect = rect_cm(target)
        for picture in slide.shapes:
            if picture.shape_type != 13:
                continue
            picture_rect = rect_cm(picture)
            if rect_contains(picture_rect, target_rect, tolerance=0.20):
                continue
            area = overlap_area(target_rect, picture_rect)
            if area > 0.01:
                issues.append(f"{shape_text(target)[:12]} 压住图片对象 {picture.name}：{area:.3f}cm²")
    return sorted(set(issues)), severe_sections


def render_check(path: Path) -> tuple[bool, str]:
    libreoffice = shutil.which("libreoffice")
    pdftotext = shutil.which("pdftotext")
    if not libreoffice or not pdftotext:
        return False, "缺少 libreoffice 或 pdftotext"
    with tempfile.TemporaryDirectory(prefix="cowork-evo-070-") as temp_dir:
        completed = subprocess.run(
            [libreoffice, "--headless", "--convert-to", "pdf", "--outdir", temp_dir, str(path)],
            capture_output=True,
            text=True,
            timeout=180,
            check=False,
        )
        pdf_path = Path(temp_dir) / f"{path.stem}.pdf"
        if completed.returncode != 0 or not pdf_path.is_file():
            return False, (completed.stderr or completed.stdout or "转换失败")[-800:]
        extracted = subprocess.run(
            [pdftotext, "-layout", str(pdf_path), "-"],
            capture_output=True,
            text=True,
            timeout=60,
            check=False,
        )
        if extracted.returncode != 0:
            return False, (extracted.stderr or "PDF 文本抽取失败")[-800:]
        pages = [page for page in extracted.stdout.split("\f") if page.strip()]
        normalized = re.sub(r"\s+", "", "".join(pages))
    missing = []
    for _, title, body_prefix, _ in SECTIONS:
        for text in (title, body_prefix):
            if re.sub(r"\s+", "", text) not in normalized:
                missing.append(text)
    return len(pages) == 1 and not missing, f"页数={len(pages)}，缺失文本={missing}"


def evaluate(artifact_dir: Path) -> dict:
    output_path = artifact_dir / OUTPUT_NAME
    results = []
    if not output_path.is_file():
        results.append(criterion("R001", False, f"缺少 {OUTPUT_NAME}"))
        return build_result(results, output_path)

    try:
        candidate = Presentation(output_path)
    except Exception as exc:
        results.append(criterion("R001", False, f"文件无法打开：{exc}"))
        return build_result(results, output_path)

    width_cm = candidate.slide_width / EMU_PER_CM
    height_cm = candidate.slide_height / EMU_PER_CM
    dim1_ok = (
        len(candidate.slides) == 1
        and abs(width_cm - 20.00) <= 0.05
        and abs(height_cm - 28.29) <= 0.05
    )
    results.append(
        criterion(
            "R001",
            dim1_ok,
            f"页数={len(candidate.slides)}，尺寸={width_cm:.4f}×{height_cm:.4f}cm",
        )
    )
    if not dim1_ok:
        return build_result(results, output_path)

    upstream = upstream_verifier._run_evaluation(str(output_path))
    hit_codes = {hit.code for hit in upstream.hits}
    for criterion_id, section in zip(("R002", "R003", "R004", "R005", "R006", "R007"), SECTIONS):
        results.append(criterion(criterion_id, section[3] in hit_codes, f"上游规则={section[3]}"))
    results.append(criterion("R008", "+3#font" in hit_codes, "上游字体规则"))
    spacing_ok = "+3#spacing" in hit_codes and "+1#align" in hit_codes
    results.append(criterion("R009", spacing_ok, f"命中规则={sorted(hit_codes)}"))

    geometry_issues, severe_sections = geometry_check(candidate.slides[0])
    results.append(
        criterion(
            "R010",
            not geometry_issues,
            "未发现目标文字重叠" if not geometry_issues else "；".join(geometry_issues[:20]),
        )
    )

    source_path = Path(__file__).resolve().parents[1] / "data" / "input_files" / INPUT_NAME
    preservation_errors = []
    try:
        source = Presentation(source_path)
        preservation_errors = compare_preserved_objects(source, candidate)
    except Exception as exc:
        preservation_errors = [f"无法加载或比较原始输入：{exc}"]
    results.append(
        criterion(
            "R011",
            not preservation_errors,
            "非目标对象完整保留" if not preservation_errors else "；".join(preservation_errors[:20]),
        )
    )

    rendered, render_evidence = render_check(output_path)
    results.append(criterion("R012", rendered, render_evidence))
    results.append(criterion("R013", severe_sections >= 3, f"问题板块数={severe_sections}"))
    results.append(criterion("R014", len(preservation_errors) >= 2, f"非目标对象错误数={len(preservation_errors)}"))
    return build_result(results, output_path)


def build_result(results: list[dict], artifact_path: Path) -> dict:
    present = {item["criterion_id"] for item in results}
    for identifier, (_, _, kind) in CRITERION_SPECS.items():
        if identifier not in present:
            results.append(criterion(identifier, False, "因交付文件不存在、损坏或门槛失败，本判据未执行"))
    results.sort(key=lambda item: item["criterion_id"])

    positive_total = sum(item["weight"] for item in results if item["weight"] > 0)
    earned = sum(item["weight"] for item in results if item["weight"] > 0 and item["passed"])
    penalties = sum(item["weight"] for item in results if item["weight"] < 0 and item["passed"])
    reward = max(0.0, min(1.0, (earned + penalties) / positive_total))
    hurdle_failures = [item["criterion_id"] for item in results if item["type"] == "hurdle" and not item["passed"]]
    if hurdle_failures:
        reward = 0.0

    if artifact_path.is_file():
        artifact_digest = hashlib.sha256(artifact_path.read_bytes()).hexdigest()
    else:
        artifact_digest = hashlib.sha256(b"").hexdigest()

    normalized = []
    for item in results:
        score = (0.0 if item["passed"] else 1.0) if item["type"] == "penalty" else item["score"]
        normalized.append(
            {
                "criterion_id": item["criterion_id"],
                "score": score,
                "evidence": item["evidence"],
                "method": item["method"],
                "raw": {
                    "description": item["description"],
                    "weight": item["weight"],
                    "type": item["type"],
                    "passed": item["passed"],
                },
            }
        )
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [artifact_digest],
        "criterion_results": normalized,
        "reward": round(reward, 6),
        "passed": reward >= 0.7,
        "hurdle_failures": hurdle_failures,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path)
    args = parser.parse_args()
    result = evaluate(args.artifact_dir)
    rendered = json.dumps(result, ensure_ascii=False, indent=2)
    print(rendered)
    if args.result:
        args.result.parent.mkdir(parents=True, exist_ok=True)
        args.result.write_text(rendered + "\n", encoding="utf-8")
    raise SystemExit(0 if result["passed"] else 1)


if __name__ == "__main__":
    main()
