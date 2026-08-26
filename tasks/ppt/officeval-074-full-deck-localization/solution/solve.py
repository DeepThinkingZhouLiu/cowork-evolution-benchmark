#!/usr/bin/env python3
"""逐页重放 officeval_074 的人工英文化参考解法。"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from translations import TRANSLATIONS


SOURCE_NAME = "农业县低空经济发展方案.pptx"
OUTPUT_NAME = "农业县低空经济发展方案_English.pptx"
CJK_RE = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff]", flags=re.UNICODE)


def translate_deck(input_dir: Path, output_dir: Path) -> Path:
    source_path = input_dir / SOURCE_NAME
    presentation = Presentation(source_path)
    if len(presentation.slides) != 46:
        raise ValueError(f"源 PPT 应为 46 页，实际 {len(presentation.slides)} 页")

    seen: set[tuple[int, int]] = set()
    for page, slide in enumerate(presentation.slides, start=1):
        for shape_index, shape in enumerate(slide.shapes):
            if not getattr(shape, "has_text_frame", False):
                continue
            source_text = shape.text_frame.text or ""
            if not CJK_RE.search(source_text):
                continue
            key = (page, shape_index)
            if key not in TRANSLATIONS:
                raise KeyError(f"缺少逐对象英译：{key} {source_text!r}")
            runs = [
                run
                for paragraph in shape.text_frame.paragraphs
                for run in paragraph.runs
                if run.text
            ]
            if not runs:
                raise ValueError(f"含中文的对象没有可替换 run：{key}")
            runs[0].text = TRANSLATIONS[key]
            for run in runs[1:]:
                run.text = ""
            # 源第 27 页该承接句首 run 异常为 36pt；恢复该内容框的既定 9.2pt。
            if key == (27, 16):
                runs[0].font.size = Pt(9.2)
            # 源文件有 7 个承接句文本框只有 0.01–0.18 英寸高，其中第 27 页还残留
            # 25% 的自动缩放参数；源中文和替换后的英文都会缩成不可读的小字。
            # 仅在原承接卡片内部重新分配标题与正文的垂直空间，不改卡片、图片或背景。
            if key in {
                (13, 16),
                (20, 16),
                (26, 16),
                (27, 16),
                (30, 16),
                (31, 16),
                (43, 16),
            }:
                transition_box = slide.shapes[14]
                transition_title = slide.shapes[15]
                shape.top = transition_title.top + Inches(0.24)
                shape.height = (
                    transition_box.top
                    + transition_box.height
                    - shape.top
                    - Inches(0.04)
                )
                shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            seen.add(key)

    missing = set(TRANSLATIONS) - seen
    if missing:
        raise KeyError(f"英译表存在未匹配对象：{sorted(missing)}")

    remaining = []
    for page, slide in enumerate(presentation.slides, start=1):
        for shape_index, shape in enumerate(slide.shapes):
            if getattr(shape, "has_text_frame", False) and CJK_RE.search(shape.text_frame.text or ""):
                remaining.append((page, shape_index, shape.text_frame.text))
    if remaining:
        raise AssertionError(f"仍有中文文本：{remaining[:10]}")

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / OUTPUT_NAME
    presentation.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    output = translate_deck(args.input_dir, args.output_dir)
    print(output)


if __name__ == "__main__":
    main()
