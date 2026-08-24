#!/usr/bin/env python3
"""officeval_075 的逐题参考解法。

本文件只处理“医疗服务收费指引.pptx”和“呼吸项目课件.pptx”这两个真实输入。
它从 5 页参考 Deck 读取角色背景与主配色，再逐页修改 39 页课件的背景、字体、
品牌脚手架和页面角色样式；不生成新题，也不改医疗正文事实。
"""

from __future__ import annotations

import argparse
import re
from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


REFERENCE_NAME = "医疗服务收费指引.pptx"
SOURCE_NAME = "呼吸项目课件.pptx"
OUTPUT_NAME = "呼吸项目收费解读_合规模板风格版.pptx"

FONT_CN = "微软雅黑"
FONT_LATIN = "Arial"
BLACK = "000000"
WHITE = "FFFFFF"
GRAY = "5C6F7A"
LIGHT_GREEN = "C6F0DC"


def rgb_distance(left: str, right: str) -> int:
    return sum(abs(int(left[i : i + 2], 16) - int(right[i : i + 2], 16)) for i in (0, 2, 4))


def explicit_fill_rgb(shape) -> str | None:
    try:
        value = str(shape.fill.fore_color.rgb)
        return value if re.fullmatch(r"[0-9A-Fa-f]{6}", value) else None
    except Exception:
        return None


def reference_style(reference: Presentation) -> tuple[list[str], dict[str, str]]:
    if len(reference.slides) != 5:
        raise ValueError(f"参考模板应为 5 页，实际 {len(reference.slides)} 页")
    backgrounds = []
    for index, slide in enumerate(reference.slides, start=1):
        try:
            color = str(slide.background.fill.fore_color.rgb)
        except Exception as exc:
            raise ValueError(f"参考模板第 {index} 页没有显式 RGB 背景：{exc}") from exc
        backgrounds.append(color)
    expected_backgrounds = ["F4FBFA", "FFFDF7", "143642", "F8FBF7", "F5FBFA"]
    if backgrounds != expected_backgrounds:
        raise ValueError(f"参考模板页面角色背景发生变化：{backgrounds}")

    fills = Counter(
        color
        for slide in reference.slides
        for shape in slide.shapes
        if (color := explicit_fill_rgb(shape)) is not None and color != WHITE
    )
    candidates = list(fills)
    if not candidates:
        raise ValueError("参考模板没有可读取的显式形状配色")

    def nearest(target: str) -> str:
        return min(candidates, key=lambda color: (rgb_distance(color, target), -fills[color]))

    palette = {
        "green": nearest("149E9A"),
        "orange": nearest("F4A635"),
        "pale": nearest("B7DCD7"),
        "dark": backgrounds[2],
    }
    if rgb_distance(palette["green"], "149E9A") > 30:
        raise ValueError(f"参考模板主绿色异常：{palette['green']}")
    if rgb_distance(palette["orange"], "F4A635") > 30:
        raise ValueError(f"参考模板强调橙异常：{palette['orange']}")
    return backgrounds, palette


def find_name(slide, name: str):
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    raise ValueError(f"第 {slide.slide_id} 页缺少对象 {name}")


def find_text(slide, text: str):
    matches = [
        shape
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and text in (shape.text_frame.text or "")
    ]
    if len(matches) != 1:
        raise ValueError(f"文本 {text!r} 应唯一匹配，实际 {len(matches)} 个")
    return matches[0]


def remove_shape(shape) -> None:
    parent = shape._element.getparent()
    if parent is not None:
        parent.remove(shape._element)


def bring_to_front(shape) -> None:
    parent = shape._element.getparent()
    if parent is None:
        return
    parent.remove(shape._element)
    parent.append(shape._element)


def rename_shape(shape, name: str) -> None:
    shape._element.nvSpPr.cNvPr.set("name", name)


def set_preset(shape, preset: str) -> None:
    shape._element.spPr.prstGeom.set("prst", preset)


def set_fill(shape, color: str) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)


def set_line(shape, color: str, width: float = 0.8) -> None:
    shape.line.color.rgb = RGBColor.from_string(color)
    shape.line.width = Pt(width)


def set_background(slide, color: str) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(color)


def style_runs(
    shape,
    *,
    font: str | None = None,
    size: float | None = None,
    bold: bool | None = None,
    color: str | None = None,
) -> None:
    if not getattr(shape, "has_text_frame", False):
        return
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if font is not None:
                run.font.name = font
            if size is not None:
                run.font.size = Pt(size)
            if bold is not None:
                run.font.bold = bold
            if color is not None:
                run.font.color.rgb = RGBColor.from_string(color)


def replace_text(shape, text: str) -> None:
    shape.text_frame.clear()
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.add_run().text = text


def style_text(
    shape,
    *,
    font: str,
    size: float,
    bold: bool,
    color: str,
    alignment: PP_ALIGN | None = None,
) -> None:
    style_runs(shape, font=font, size=size, bold=bold, color=color)
    if alignment is not None:
        for paragraph in shape.text_frame.paragraphs:
            paragraph.alignment = alignment


def remove_decorative_ovals(slide) -> None:
    for name in ("Oval 3", "Oval 4", "Oval 5"):
        for shape in list(slide.shapes):
            if shape.name == name:
                remove_shape(shape)


def recolor_existing_shapes(presentation: Presentation, palette: dict[str, str]) -> None:
    mapping = {
        "17979D": palette["green"],
        "4EBFB8": palette["pale"],
        "E3AA41": palette["orange"],
        "133545": palette["dark"],
        "005A64": palette["dark"],
        "BDD3D2": palette["pale"],
        "DAEEEB": "DDEEEA",
    }
    for slide in presentation.slides:
        for shape in slide.shapes:
            old = explicit_fill_rgb(shape)
            if old in mapping:
                set_fill(shape, mapping[old])


def set_global_scaffold(
    presentation: Presentation,
    backgrounds: list[str],
    palette: dict[str, str],
) -> None:
    if len(presentation.slides) != 39:
        raise ValueError(f"待改课件应为 39 页，实际 {len(presentation.slides)} 页")
    role_background = {
        1: backgrounds[0],
        2: backgrounds[1],
        3: backgrounds[2],
        8: backgrounds[2],
        39: backgrounds[4],
    }
    for page, slide in enumerate(presentation.slides, start=1):
        set_background(slide, role_background.get(page, backgrounds[3]))
        remove_decorative_ovals(slide)

        top_bar = find_name(slide, "Rectangle 1")
        top_bar.left = 0
        top_bar.top = 0
        top_bar.width = presentation.slide_width
        top_bar.height = Inches(0.14)
        set_fill(top_bar, palette["green"])
        set_line(top_bar, palette["green"], 0.1)

        bottom_rule = find_name(slide, "Rectangle 2")
        bottom_rule.left = Inches(0.55)
        bottom_rule.top = Inches(7.02)
        bottom_rule.width = Inches(12.23)
        bottom_rule.height = Inches(0.018)
        set_fill(bottom_rule, palette["green"])
        set_line(bottom_rule, palette["green"], 0.1)

        for shape in slide.shapes:
            style_runs(shape, font=FONT_CN)

        footer = find_name(slide, "TextBox 6")
        footer.left = Inches(0.62)
        footer.top = Inches(7.13)
        footer.width = Inches(5.3)
        footer.height = Inches(0.18)
        style_text(footer, font=FONT_CN, size=5.5, bold=False, color="737373")

        page_number = find_name(slide, "TextBox 7")
        replace_text(page_number, f"{page:02d}/39")
        page_number.left = Inches(11.75)
        page_number.top = Inches(7.08)
        page_number.width = Inches(1.02)
        page_number.height = Inches(0.24)
        style_text(
            page_number,
            font=FONT_CN,
            size=8,
            bold=True,
            color=palette["green"],
            alignment=PP_ALIGN.RIGHT,
        )
        bring_to_front(footer)
        bring_to_front(page_number)


def brand_chip(slide, palette: dict[str, str], *, existing=None) -> None:
    chip = existing
    if chip is None:
        chip = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.62),
            Inches(0.18),
            Inches(2.35),
            Inches(0.34),
        )
    replace_text(chip, "医疗价格服务价格管理")
    set_preset(chip, "roundRect")
    set_fill(chip, palette["green"])
    set_line(chip, palette["green"], 0.4)
    chip.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    chip.text_frame.margin_left = Inches(0.05)
    chip.text_frame.margin_right = Inches(0.05)
    chip.text_frame.margin_top = 0
    chip.text_frame.margin_bottom = 0
    style_text(
        chip,
        font=FONT_CN,
        size=10.5,
        bold=True,
        color=WHITE,
        alignment=PP_ALIGN.CENTER,
    )


def crop_picture_to_box(picture, *, left, top, width, height) -> None:
    picture.left = left
    picture.top = top
    picture.width = width
    picture.height = height
    picture.crop_top = 0
    picture.crop_bottom = 0
    picture.crop_left = 0.24
    picture.crop_right = 0.24


def style_cover(slide, palette: dict[str, str]) -> None:
    brand_chip(slide, palette, existing=find_name(slide, "Rounded Rectangle 10"))
    chip = find_name(slide, "Rounded Rectangle 10")
    chip.left = Inches(0.78)
    chip.top = Inches(0.62)
    chip.width = Inches(2.4)
    chip.height = Inches(0.38)

    picture = find_name(slide, "Picture 8")
    crop_picture_to_box(
        picture,
        left=Inches(7.0),
        top=Inches(0.62),
        width=Inches(5.7),
        height=Inches(5.95),
    )
    remove_shape(find_name(slide, "Rectangle 9"))

    # 原封面上的“医疗服务价格管理”是实质内容，不能因模板品牌条而丢失。
    # 将它作为图片区的小标签恢复，同时保留模板要求的品牌条。
    service_label = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(9.55),
        Inches(0.82),
        Inches(2.35),
        Inches(0.38),
    )
    replace_text(service_label, "医疗服务价格管理")
    set_fill(service_label, palette["dark"])
    set_line(service_label, palette["dark"], 0.4)
    service_label.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    style_text(
        service_label,
        font=FONT_CN,
        size=10.5,
        bold=True,
        color=WHITE,
        alignment=PP_ALIGN.CENTER,
    )
    bring_to_front(service_label)

    for name in ("Arc 15", "Arc 16"):
        remove_shape(find_name(slide, name))
    accent_line = find_name(slide, "Rectangle 17")
    accent_line.left = Inches(0.9)
    accent_line.top = Inches(5.34)
    accent_line.width = Inches(4.25)
    accent_line.height = Inches(0.025)
    set_fill(accent_line, palette["green"])

    title = find_name(slide, "TextBox 11")
    replace_text(title, "呼吸系统诊疗收费项目合规解读")
    title.left = Inches(0.9)
    title.top = Inches(1.55)
    title.width = Inches(5.25)
    title.height = Inches(1.2)
    title.text_frame.word_wrap = True
    style_text(title, font=FONT_CN, size=34, bold=True, color=BLACK)

    subtitle = find_name(slide, "TextBox 12")
    subtitle.left = Inches(0.9)
    subtitle.top = Inches(3.2)
    subtitle.width = Inches(5.5)
    subtitle.height = Inches(0.55)
    style_text(subtitle, font=FONT_CN, size=15.5, bold=False, color=GRAY)

    organization = find_name(slide, "TextBox 13")
    organization.left = Inches(0.9)
    organization.top = Inches(4.25)
    organization.width = Inches(2.25)
    organization.height = Inches(0.39)
    set_fill(organization, palette["green"])
    organization.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    style_text(
        organization,
        font=FONT_CN,
        size=9.5,
        bold=True,
        color=WHITE,
        alignment=PP_ALIGN.CENTER,
    )

    date = find_name(slide, "TextBox 14")
    date.left = Inches(3.32)
    date.top = Inches(4.25)
    date.width = Inches(1.65)
    date.height = Inches(0.39)
    set_fill(date, palette["orange"])
    date.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    style_text(date, font=FONT_CN, size=9.5, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)


def style_directory(slide, palette: dict[str, str]) -> None:
    brand_chip(slide, palette)
    style_text(find_name(slide, "TextBox 8"), font=FONT_CN, size=28, bold=True, color=BLACK)
    style_text(find_name(slide, "TextBox 9"), font=FONT_CN, size=12.5, bold=False, color=GRAY)

    specs = (
        ("Oval 13", "TextBox 14", "TextBox 15", "01", palette["green"]),
        ("Oval 18", "TextBox 19", "TextBox 20", "02", palette["orange"]),
        ("Oval 23", "TextBox 24", "TextBox 25", "03", palette["green"]),
    )
    for oval_name, title_name, body_name, number, fill in specs:
        oval = find_name(slide, oval_name)
        original_left, original_top, original_width, original_height = (
            oval.left,
            oval.top,
            oval.width,
            oval.height,
        )
        replace_text(oval, "")
        set_fill(oval, fill)
        set_line(oval, fill, 0.4)

        number_box = slide.shapes.add_textbox(
            original_left + Inches(0.04),
            original_top + Inches(0.04),
            original_width - Inches(0.08),
            Inches(0.22),
        )
        rename_shape(number_box, oval_name)
        replace_text(number_box, number)
        number_box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        number_box.text_frame.margin_left = 0
        number_box.text_frame.margin_right = 0
        number_box.text_frame.margin_top = 0
        number_box.text_frame.margin_bottom = 0
        style_text(
            number_box,
            font=FONT_LATIN,
            size=10.5,
            bold=True,
            color=WHITE,
            alignment=PP_ALIGN.CENTER,
        )
        style_text(find_name(slide, title_name), font=FONT_CN, size=15, bold=True, color=BLACK)
        style_text(find_name(slide, body_name), font=FONT_CN, size=9.8, bold=False, color=GRAY)


def style_part(slide, palette: dict[str, str], *, part_number: int) -> None:
    picture = find_name(slide, "Picture 8")
    crop_picture_to_box(
        picture,
        left=Inches(7.05),
        top=Inches(0.72),
        width=Inches(5.35),
        height=Inches(5.95),
    )
    panel = find_name(slide, "Rectangle 9")
    panel.left = 0
    panel.top = 0
    panel.width = Inches(7.25)
    panel.height = Inches(7.5)
    set_fill(panel, palette["dark"])
    set_line(panel, palette["dark"], 0.1)

    part = find_name(slide, "TextBox 10")
    set_preset(part, "roundRect")
    replace_text(part, f"part {part_number:02d}")
    part.left = Inches(0.85)
    part.top = Inches(1.0)
    part.width = Inches(1.35)
    part.height = Inches(0.42)
    set_fill(part, palette["green"])
    set_line(part, palette["green"], 0.4)
    part.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    style_text(part, font=FONT_LATIN, size=10.5, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)

    heading = find_name(slide, "TextBox 11")
    heading.left = Inches(0.88)
    heading.top = Inches(2.1)
    heading.width = Inches(5.7)
    heading.height = Inches(1.0)
    style_text(heading, font=FONT_CN, size=34, bold=True, color=WHITE)

    subtitle = find_name(slide, "TextBox 12")
    subtitle.left = Inches(0.92)
    subtitle.top = Inches(3.55)
    subtitle.width = Inches(5.8)
    subtitle.height = Inches(0.65)
    style_text(subtitle, font=FONT_CN, size=14, bold=False, color=LIGHT_GREEN)


def style_number(shape) -> None:
    style_text(shape, font=FONT_LATIN, size=10, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)


def style_policy_pages(presentation: Presentation, palette: dict[str, str]) -> None:
    for page in range(4, 8):
        slide = presentation.slides[page - 1]
        style_text(find_name(slide, "TextBox 8"), font=FONT_CN, size=27, bold=True, color=BLACK)
        style_text(find_name(slide, "TextBox 9"), font=FONT_CN, size=11.5, bold=False, color=GRAY)
        for title_name in ("TextBox 14", "TextBox 19", "TextBox 24"):
            style_text(find_name(slide, title_name), font=FONT_CN, size=14.5, bold=True, color=BLACK)
        for body_name in ("TextBox 15", "TextBox 20", "TextBox 25"):
            style_text(find_name(slide, body_name), font=FONT_CN, size=9.7, bold=False, color=GRAY)
        for number_name in ("Oval 13", "Oval 18", "Oval 23"):
            style_number(find_name(slide, number_name))

        # 第 4 页比其他政策页多两个横向说明区，也要纳入同一层级。
        # 源文件的最后一个正文框高度过小，LibreOffice 会渲染成难以阅读的小字。
        if page == 4:
            for title_name in ("TextBox 28", "TextBox 32"):
                style_text(
                    find_name(slide, title_name),
                    font=FONT_CN,
                    size=12.5,
                    bold=True,
                    color=palette["dark"],
                )
            for body_name in ("TextBox 29", "TextBox 33"):
                style_text(
                    find_name(slide, body_name),
                    font=FONT_CN,
                    size=9.7,
                    bold=False,
                    color=GRAY,
                )
            find_name(slide, "TextBox 33").height = Inches(0.42)


def style_detail_pages(presentation: Presentation, palette: dict[str, str]) -> None:
    for page in range(9, 38):
        slide = presentation.slides[page - 1]
        style_text(find_name(slide, "TextBox 8"), font=FONT_CN, size=27, bold=True, color=BLACK)
        style_text(find_name(slide, "TextBox 9"), font=FONT_CN, size=11.5, bold=False, color=GRAY)

        tag = find_name(slide, "Rounded Rectangle 11")
        set_fill(tag, palette["green"])
        set_line(tag, palette["green"], 0.4)
        style_text(tag, font=FONT_CN, size=10.5, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
        style_text(find_name(slide, "TextBox 12"), font=FONT_CN, size=9.7, bold=False, color=GRAY)

        for title_name in ("TextBox 16", "TextBox 22", "TextBox 28"):
            style_text(find_name(slide, title_name), font=FONT_CN, size=14.5, bold=True, color=BLACK)
        for body_name in ("TextBox 17", "TextBox 23", "TextBox 29"):
            style_text(find_name(slide, body_name), font=FONT_CN, size=9.7, bold=False, color=GRAY)

        # 每个明细页底部还有一个“收费审核重点”说明区。
        # 源文本框高度只有 0.08 英寸，必须在保留原文的前提下恢复可读性。
        style_text(
            find_name(slide, "TextBox 32"),
            font=FONT_CN,
            size=12.5,
            bold=True,
            color=palette["dark"],
        )
        detail_note = find_name(slide, "TextBox 33")
        style_text(detail_note, font=FONT_CN, size=9.2, bold=False, color=GRAY)
        detail_note.height = Inches(0.42)

        for old_name, allowed_name in (("Oval 15", "Oval 13"), ("Oval 21", "Oval 18"), ("Oval 27", "Oval 23")):
            number = find_name(slide, old_name)
            rename_shape(number, allowed_name)
            style_number(number)


def style_attention_page(slide, palette: dict[str, str]) -> None:
    style_text(find_name(slide, "TextBox 8"), font=FONT_CN, size=27, bold=True, color=BLACK)
    style_text(find_name(slide, "TextBox 9"), font=FONT_CN, size=11.5, bold=False, color=GRAY)
    for title_name in ("TextBox 14", "TextBox 19", "TextBox 24"):
        style_text(find_name(slide, title_name), font=FONT_CN, size=14.5, bold=True, color=BLACK)
    for body_name in ("TextBox 15", "TextBox 20", "TextBox 25"):
        style_text(find_name(slide, body_name), font=FONT_CN, size=9.7, bold=False, color=GRAY)
    for number_name in ("Oval 13", "Oval 18", "Oval 23"):
        style_number(find_name(slide, number_name))
    tag = find_name(slide, "Rounded Rectangle 26")
    set_fill(tag, palette["green"])
    set_line(tag, palette["green"], 0.4)
    style_text(tag, font=FONT_CN, size=10.5, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)


def style_end(slide, palette: dict[str, str]) -> None:
    picture = find_name(slide, "Picture 8")
    crop_picture_to_box(
        picture,
        left=Inches(7.15),
        top=Inches(0.82),
        width=Inches(5.3),
        height=Inches(5.85),
    )
    card = find_name(slide, "Rectangle 9")
    set_preset(card, "roundRect")
    card.left = Inches(0.78)
    card.top = Inches(1.15)
    card.width = Inches(5.85)
    card.height = Inches(5.25)
    set_fill(card, WHITE)
    set_line(card, palette["pale"], 0.8)

    title = find_name(slide, "TextBox 10")
    title.left = Inches(1.55)
    title.top = Inches(2.15)
    title.width = Inches(4.3)
    title.height = Inches(0.8)
    style_text(title, font=FONT_CN, size=36, bold=True, color=BLACK, alignment=PP_ALIGN.CENTER)

    subtitle = find_name(slide, "TextBox 11")
    subtitle.left = Inches(1.3)
    subtitle.top = Inches(3.35)
    subtitle.width = Inches(4.85)
    subtitle.height = Inches(0.6)
    style_text(subtitle, font=FONT_CN, size=12, bold=False, color=GRAY, alignment=PP_ALIGN.CENTER)

    organization = find_name(slide, "TextBox 12")
    organization.left = Inches(1.7)
    organization.top = Inches(4.45)
    organization.width = Inches(4.05)
    organization.height = Inches(0.45)
    style_text(
        organization,
        font=FONT_CN,
        size=13.5,
        bold=True,
        color=palette["green"],
        alignment=PP_ALIGN.CENTER,
    )


def solve(input_dir: Path, output_dir: Path) -> Path:
    reference = Presentation(input_dir / REFERENCE_NAME)
    presentation = Presentation(input_dir / SOURCE_NAME)
    backgrounds, palette = reference_style(reference)

    recolor_existing_shapes(presentation, palette)
    set_global_scaffold(presentation, backgrounds, palette)
    style_cover(presentation.slides[0], palette)
    style_directory(presentation.slides[1], palette)
    style_part(presentation.slides[2], palette, part_number=1)
    style_part(presentation.slides[7], palette, part_number=2)

    for page, slide in enumerate(presentation.slides, start=1):
        if page not in (3, 8, 39):
            if page != 1:
                brand_chip(slide, palette)

    style_policy_pages(presentation, palette)
    style_detail_pages(presentation, palette)
    style_attention_page(presentation.slides[37], palette)
    style_end(presentation.slides[38], palette)

    # 新增/移动对象之后再次将页脚与页码置于最上层，保证深色或图片页也可读。
    for slide in presentation.slides:
        bring_to_front(find_name(slide, "TextBox 6"))
        bring_to_front(find_name(slide, "TextBox 7"))

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / OUTPUT_NAME
    presentation.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=Path("/app/data/input_files"))
    parser.add_argument("--output-dir", type=Path, default=Path("/app/output"))
    args = parser.parse_args()
    solve(args.input_dir, args.output_dir)


if __name__ == "__main__":
    main()
