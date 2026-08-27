#!/usr/bin/env python3
"""officeval_075 的 Harbor Judge。

上游 Verifier 原样保留在同目录。本 Judge 直接复用其正确规则，
并以真实输入文件修正 H/J/L 三处与源文件冲突的判定。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import subprocess
import tempfile
from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import upstream_verifier


OUTPUT_NAME = "呼吸项目收费解读_合规模板风格版.pptx"
SOURCE_NAME = "呼吸项目课件.pptx"
TASK_ID = "cowork-evo/officeval-075-template-style-transfer"
RUBRIC_VERSION = "omegause-officeval-075-adaptation-v0.1"
JUDGE_VERSION = "officeval-075-adaptation-v0.1"
PASS_THRESHOLD = 0.8

CRITERION_SPECS = {
    "R001": ("指定 PPTX 可打开且保持 39 页", 2, "hurdle"),
    "R002": ("原始文字、图片和页序保持，不用整页截图替代", 3, "hurdle"),
    "R003": ("39 页背景按参考 Deck 的页面角色迁移", 5, "positive"),
    "R004": ("全稿字体遵循微软雅黑与编号例外", 5, "positive"),
    "R005": ("页脚和 01/39–39/39 页码完整正确", 10, "positive"),
    "R006": ("封面文字层级、主视觉和强调色正确", 1, "positive"),
    "R007": ("品牌标识和两张章节页标识正确", 6, "positive"),
    "R008": ("目录页和章节页样式正确，保留源文引号", 2, "positive"),
    "R009": ("第 4–7 页政策页按原竖向卡片角色正确排版", 3, "positive"),
    "R010": ("第 9–37 页项目明细页层级正确", 5, "positive"),
    "R011": ("第 38 页保留“不得同时收费”并完成指定样式", 1, "positive"),
    "R012": ("编号样式和结束页完成度正确", 6, "positive"),
    "R013": ("LibreOffice 实际渲染为 39 页且关键文字可读", 2, "positive"),
    "R014": ("原文、原图丢失、页序损坏或新增截图压平", -12, "penalty"),
    "R015": ("出现批注、临时标记、异常嵌入物或大量多余对象", -6, "penalty"),
}


def criterion(identifier: str, passed: bool, evidence: str = "") -> dict:
    description, weight, kind = CRITERION_SPECS[identifier]
    return {
        "criterion_id": identifier,
        "description": description,
        "weight": weight,
        "type": kind,
        "passed": bool(passed),
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
    }


def normalize(text: str) -> str:
    return re.sub(r"\s+", "", text or "")


def upstream_results(candidate_path: Path) -> dict[str, object]:
    return {item.rule_id: item for item in upstream_verifier.check_dim2(str(candidate_path))}


def upstream_pair(results: dict[str, object], *identifiers: str) -> tuple[bool, str]:
    selected = [results[identifier] for identifier in identifiers]
    return all(item.hit for item in selected), " | ".join(
        f"{item.rule_id}: {item.detail}" for item in selected
    )


def semantic_texts(slide) -> Counter:
    """比较业务文字，仅排除允许随模板变化的页码和脚手标签。"""
    values = []
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        value = normalize(shape.text_frame.text)
        if not value or shape.name == "TextBox 7":
            continue
        if value in {"医疗服务价格管理", "医疗价格服务价格管理"}:
            continue
        if re.fullmatch(r"(?i)part0?[12]", value):
            continue
        values.append(value)
    return Counter(values)


def iter_pictures(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            yield shape
        elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_pictures(shape.shapes)


def picture_hashes(slide) -> Counter:
    return Counter(hashlib.sha256(shape.image.blob).hexdigest() for shape in iter_pictures(slide.shapes))


def preservation_check(source: Presentation, candidate: Presentation) -> tuple[bool, list[str]]:
    errors = []
    if len(source.slides) != len(candidate.slides):
        return False, [f"页数 {len(source.slides)} -> {len(candidate.slides)}"]
    for index, (before, after) in enumerate(zip(source.slides, candidate.slides), start=1):
        if semantic_texts(before) != semantic_texts(after):
            missing = semantic_texts(before) - semantic_texts(after)
            added = semantic_texts(after) - semantic_texts(before)
            errors.append(f"p{index} 文字差异：缺失={list(missing.elements())[:4]}，新增={list(added.elements())[:4]}")
        if picture_hashes(before) != picture_hashes(after):
            missing = picture_hashes(before) - picture_hashes(after)
            added = picture_hashes(after) - picture_hashes(before)
            errors.append(f"p{index} 图片差异：缺失={sum(missing.values())}，新增={sum(added.values())}")
    return not errors, errors


def shape_by_name(slide, name: str):
    for shape in slide.shapes:
        if shape.name == name and getattr(shape, "has_text_frame", False):
            return shape
    return None


def exact_text_shape(slide, text: str):
    target = normalize(text)
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and normalize(shape.text_frame.text) == target:
            return shape
    return None


def style_check(
    slide,
    shape_name: str,
    text: str,
    *,
    size: float,
    color: str,
    bold: bool,
    font: str = "微软雅黑",
) -> tuple[bool, str]:
    shape = shape_by_name(slide, shape_name)
    if shape is None:
        return False, f"缺少 {shape_name}"
    if normalize(shape.text_frame.text) != normalize(text):
        return False, f"{shape_name} 文字={shape.text_frame.text!r}"
    ok, detail = upstream_verifier._check_font_attr(
        list(upstream_verifier._iter_runs(shape)),
        text,
        size=size,
        color=color,
        bold=bold,
        font_name=font,
        size_tol=0.6,
        color_tol=22,
    )
    return ok, f"{shape_name}: {detail}"


def text_style_check(
    slide,
    text: str,
    *,
    size: float,
    color: str,
    bold: bool,
    font: str = "微软雅黑",
) -> tuple[bool, str]:
    shape = exact_text_shape(slide, text)
    if shape is None:
        return False, f"缺少文字 {text!r}"
    actual_text = shape.text_frame.text
    ok, detail = upstream_verifier._check_font_attr(
        list(upstream_verifier._iter_runs(shape)),
        actual_text,
        size=size,
        color=color,
        bold=bold,
        font_name=font,
        size_tol=0.6,
        color_tol=22,
    )
    return ok, f"{text}: {detail}"


def combine(checks: list[tuple[bool, str]]) -> tuple[bool, str]:
    return all(item[0] for item in checks), "；".join(item[1] for item in checks)


def corrected_directory_check(candidate: Presentation) -> tuple[bool, str]:
    slide = candidate.slides[1]
    checks = [
        style_check(slide, "TextBox 8", "目录", size=28, color="000000", bold=True),
        style_check(
            slide,
            "TextBox 9",
            "从政策口径到项目明细，按“先规则、后执行”的路径展开",
            size=12.5,
            color="5C6F7A",
            bold=False,
        ),
    ]
    for name, text in (("TextBox 14", "政策文件要点"), ("TextBox 19", "项目明细规范"), ("TextBox 24", "执行注意事项")):
        checks.append(style_check(slide, name, text, size=15, color="000000", bold=True))
    for name, text in (
        ("TextBox 15", "理解整合口径、医保衔接与价格边界"),
        ("TextBox 20", "按检查、治疗、手术三类建立收费判断"),
        ("TextBox 25", "聚焦互斥规则、计价单位和复杂条件"),
    ):
        checks.append(style_check(slide, name, text, size=9.8, color="5C6F7A", bold=False))
    for number in ("01", "02", "03"):
        checks.append(text_style_check(slide, number, size=10.5, color="FFFFFF", bold=True, font="Arial"))
    return combine(checks)


def corrected_policy_check(candidate: Presentation) -> tuple[bool, str]:
    checks = []
    expected = {
        4: ("规范整合呼吸系统医疗服务项目", "本页承接政策框架，说明项目整理与医保衔接的基本方向"),
        5: ("项目类别与价格边界", "承接上一页政策概览，进一步拆解实际收费判断口径"),
        6: ("加收、扩展与基本物耗", "承接上一页政策概览，进一步拆解实际收费判断口径"),
        7: ("无创、内镜与手术并行规则", "承接上一页政策概览，进一步拆解实际收费判断口径"),
    }
    for page, (title, subtitle) in expected.items():
        slide = candidate.slides[page - 1]
        checks.append(style_check(slide, "TextBox 8", title, size=27, color="000000", bold=True))
        checks.append(style_check(slide, "TextBox 9", subtitle, size=11.5, color="5C6F7A", bold=False))
        for name in ("TextBox 14", "TextBox 19", "TextBox 24"):
            shape = shape_by_name(slide, name)
            text = shape.text_frame.text if shape is not None else ""
            checks.append(style_check(slide, name, text, size=14.5, color="000000", bold=True))
        for name in ("TextBox 15", "TextBox 20", "TextBox 25"):
            shape = shape_by_name(slide, name)
            text = shape.text_frame.text if shape is not None else ""
            checks.append(style_check(slide, name, text, size=9.7, color="5C6F7A", bold=False))
        if page == 4:
            for name in ("TextBox 28", "TextBox 32"):
                shape = shape_by_name(slide, name)
                text = shape.text_frame.text if shape is not None else ""
                checks.append(style_check(slide, name, text, size=12.5, color="143642", bold=True))
            for name in ("TextBox 29", "TextBox 33"):
                shape = shape_by_name(slide, name)
                text = shape.text_frame.text if shape is not None else ""
                checks.append(style_check(slide, name, text, size=9.7, color="5C6F7A", bold=False))
    return combine(checks)


def corrected_attention_check(candidate: Presentation) -> tuple[bool, str]:
    slide = candidate.slides[37]
    checks = [
        style_check(slide, "TextBox 8", "执行特别注意事项", size=27, color="000000", bold=True),
        style_check(slide, "TextBox 9", "承接明细页，提炼最容易出错的三类收费风险", size=11.5, color="5C6F7A", bold=False),
        style_check(slide, "TextBox 14", "不得同时收费", size=14.5, color="000000", bold=True),
        style_check(slide, "TextBox 19", "计价单位不同", size=14.5, color="000000", bold=True),
        style_check(slide, "TextBox 24", "复杂条件须满足", size=14.5, color="000000", bold=True),
        text_style_check(slide, "规范收费·留痕可查", size=10.5, color="FFFFFF", bold=True),
    ]
    return combine(checks)


def detail_extra_check(candidate: Presentation) -> tuple[bool, str]:
    checks = []
    for page in range(9, 38):
        slide = candidate.slides[page - 1]
        title = shape_by_name(slide, "TextBox 32")
        body = shape_by_name(slide, "TextBox 33")
        checks.append(
            style_check(
                slide,
                "TextBox 32",
                title.text_frame.text if title is not None else "",
                size=12.5,
                color="143642",
                bold=True,
            )
        )
        checks.append(
            style_check(
                slide,
                "TextBox 33",
                body.text_frame.text if body is not None else "",
                size=9.2,
                color="5C6F7A",
                bold=False,
            )
        )
        if body is not None:
            checks.append((body.height >= 250000, f"p{page} TextBox 33 高度={body.height}"))
    return combine(checks)


def package_penalty(path: Path, source: Presentation, candidate: Presentation) -> tuple[bool, str]:
    suspicious = []
    for page, slide in enumerate(candidate.slides, start=1):
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = shape.text_frame.text or ""
            if re.search(r"TODO|TBD|待补|占位|示意文字", text, flags=re.I):
                suspicious.append(f"p{page}:{text[:30]}")
    extra_shapes = sum(len(slide.shapes) for slide in candidate.slides) - sum(
        len(slide.shapes) for slide in source.slides
    )
    triggered = bool(suspicious) or extra_shapes > 50
    return triggered, f"临时标记={suspicious[:5]}，新增对象净值={extra_shapes}"


def render_check(path: Path) -> tuple[bool, str]:
    libreoffice = shutil.which("libreoffice")
    pdfinfo = shutil.which("pdfinfo")
    pdftotext = shutil.which("pdftotext")
    if not libreoffice or not pdfinfo or not pdftotext:
        return False, "缺少 libreoffice、pdfinfo 或 pdftotext"
    with tempfile.TemporaryDirectory(prefix="cowork-evo-075-") as temp_dir:
        converted = subprocess.run(
            [libreoffice, "--headless", "--convert-to", "pdf", "--outdir", temp_dir, str(path)],
            capture_output=True,
            text=True,
            timeout=240,
            check=False,
        )
        pdf_path = Path(temp_dir) / f"{path.stem}.pdf"
        if converted.returncode != 0 or not pdf_path.is_file():
            return False, (converted.stderr or converted.stdout or "PDF 转换失败")[-1000:]
        info = subprocess.run([pdfinfo, str(pdf_path)], capture_output=True, text=True, timeout=60, check=False)
        text_result = subprocess.run(
            [pdftotext, "-layout", str(pdf_path), "-"], capture_output=True, text=True, timeout=120, check=False
        )
        page_match = re.search(r"^Pages:\s+(\d+)", info.stdout, flags=re.M)
        pages = int(page_match.group(1)) if page_match else 0
        normalized = normalize(text_result.stdout)
    required = ["呼吸系统诊疗收费项目合规解读", "目录", "不得同时收费", "感谢聆听"]
    missing = [text for text in required if normalize(text) not in normalized]
    return pages == 39 and not missing, f"页数={pages}，缺失文字={missing}"


def evaluate(artifact_dir: Path) -> dict:
    output_path = artifact_dir / OUTPUT_NAME
    results = []
    if not output_path.is_file():
        results.append(criterion("R001", False, f"缺少 {OUTPUT_NAME}"))
        return build_result(results, output_path)
    try:
        candidate = Presentation(output_path)
    except Exception as exc:
        results.append(criterion("R001", False, f"文件无法打开：{exc}"))
        return build_result(results, output_path)

    dim1_ok, dim1_notes = upstream_verifier.check_dim1(str(output_path))
    pages_ok = dim1_ok and len(candidate.slides) == 39
    results.append(criterion("R001", pages_ok, f"{dim1_notes}；页数={len(candidate.slides)}"))

    source_path = Path(__file__).resolve().parents[1] / "data" / "input_files" / SOURCE_NAME
    try:
        source = Presentation(source_path)
        preserved, preservation_errors = preservation_check(source, candidate)
    except Exception as exc:
        source = Presentation()
        preserved, preservation_errors = False, [f"保留性比较失败：{exc}"]
    results.append(
        criterion(
            "R002",
            preserved,
            "39 页业务文字和原图逐页一致" if preserved else "；".join(preservation_errors[:20]),
        )
    )

    try:
        upstream = upstream_results(output_path)
    except Exception as exc:
        upstream = {}
        results.append(criterion("R003", False, f"上游检查无法执行：{exc}"))

    mapping = {
        "R003": ("A",),
        "R004": ("B",),
        "R005": ("C", "D"),
        "R006": ("E",),
        "R007": ("F", "G"),
        "R012": ("M", "N"),
    }
    for identifier, upstream_ids in mapping.items():
        if identifier == "R003" and any(item["criterion_id"] == "R003" for item in results):
            continue
        if not upstream:
            results.append(criterion(identifier, False, "上游检查未返回结果"))
            continue
        passed, evidence = upstream_pair(upstream, *upstream_ids)
        results.append(criterion(identifier, passed, evidence))

    directory_ok, directory_evidence = corrected_directory_check(candidate)
    part_ok, part_evidence = upstream_pair(upstream, "I") if upstream else (False, "缺少上游 I")
    results.append(criterion("R008", directory_ok and part_ok, f"目录={directory_evidence} | 章节={part_evidence}"))

    policy_ok, policy_evidence = corrected_policy_check(candidate)
    results.append(criterion("R009", policy_ok, policy_evidence))

    detail_ok, detail_evidence = upstream_pair(upstream, "K") if upstream else (False, "缺少上游 K")
    extra_ok, extra_evidence = detail_extra_check(candidate)
    results.append(criterion("R010", detail_ok and extra_ok, f"上游={detail_evidence} | 底部说明={extra_evidence}"))

    attention_ok, attention_evidence = corrected_attention_check(candidate)
    results.append(criterion("R011", attention_ok, attention_evidence))

    rendered, render_evidence = render_check(output_path)
    results.append(criterion("R013", rendered, render_evidence))
    results.append(
        criterion(
            "R014",
            not preserved or not pages_ok,
            "；".join(preservation_errors[:20]) if not preserved or not pages_ok else "未发现原文、原图、页序或截图压平问题",
        )
    )
    extra_penalty, extra_evidence = package_penalty(output_path, source, candidate)
    results.append(criterion("R015", extra_penalty, extra_evidence))
    return build_result(results, output_path)


def build_result(results: list[dict], artifact_path: Path) -> dict:
    present = {item["criterion_id"] for item in results}
    for identifier in CRITERION_SPECS:
        if identifier not in present:
            results.append(criterion(identifier, False, "因交付不存在、损坏或门槛失败，本判据未执行"))
    results.sort(key=lambda item: item["criterion_id"])

    positive_total = sum(item["weight"] for item in results if item["weight"] > 0)
    earned = sum(item["weight"] for item in results if item["weight"] > 0 and item["passed"])
    penalties = sum(item["weight"] for item in results if item["weight"] < 0 and item["passed"])
    reward = max(0.0, min(1.0, (earned + penalties) / positive_total))
    hurdle_failures = [
        item["criterion_id"] for item in results if item["type"] == "hurdle" and not item["passed"]
    ]
    if hurdle_failures:
        reward = 0.0

    digest = hashlib.sha256(artifact_path.read_bytes()).hexdigest() if artifact_path.is_file() else hashlib.sha256(b"").hexdigest()
    normalized = []
    for item in results:
        score = (0.0 if item["passed"] else 1.0) if item["type"] == "penalty" else item["score"]
        normalized.append(
            {
                "criterion_id": item["criterion_id"],
                "score": score,
                "evidence": item["evidence"],
                "method": item["method"],
                "raw": {"description": item["description"], "weight": item["weight"], "type": item["type"], "passed": item["passed"]},
            }
        )
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [digest],
        "criterion_results": normalized,
        "reward": round(reward, 6),
        "passed": reward >= PASS_THRESHOLD,
        "hurdle_failures": hurdle_failures,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path)
    args = parser.parse_args()
    result = evaluate(args.artifact_dir)
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))

    if args.reward_file:
        args.reward_file.parent.mkdir(parents=True, exist_ok=True)
        args.reward_file.write_text(f"{result['reward']:.6f}\n", encoding="utf-8")
    raise SystemExit(0 if result["passed"] else 1)


if __name__ == "__main__":
    main()
