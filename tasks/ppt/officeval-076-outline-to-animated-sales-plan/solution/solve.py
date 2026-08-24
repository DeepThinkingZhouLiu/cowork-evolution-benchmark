#!/usr/bin/env python3
"""为 OmegaUse-OfficeVal officeval_076 制作可编辑、可放映的 21 页参考交付物。"""

from __future__ import annotations

import argparse
import hashlib
import json
import tempfile
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


OUTPUT_NAME = "艺术漆北区销售额提升方案.pptx"
TEMPLATE_NAME = "模版.pptx"
TEMPLATE_SHA256 = "fae5afedb2b42e6b0c4d729f0e99682fbc041b95cd0dfbcb04b3c6a3c0cd0a8b"
TEMPLATE_BACKGROUND_SHA256 = "9112d4e712200684e80dc40d0dea06c6e5747f43f08e41b1bf7560331eb729c6"

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NSMAP = {"p": P_NS, "a": A_NS, "r": R_NS}

FONT_NAME = "宋体"
TEXT_DARK = RGBColor(0x17, 0x37, 0x2C)
TEXT_MUTED = RGBColor(0x35, 0x5A, 0x49)
ORANGE = RGBColor(0xE6, 0x7E, 0x22)
ORANGE_TEXT = RGBColor(0xA9, 0x45, 0x00)
GREEN_LIGHT = "EAF7EE"
GREEN_MEDIUM = "7FC99A"


def sha256(path: Path) -> str:
    """计算文件的 SHA256。"""
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def set_east_asian_font(run, name: str = FONT_NAME) -> None:
    """同时写入拉丁与东亚字体，避免 Office 回退到主题字体。"""
    run.font.name = name
    r_pr = run._r.get_or_add_rPr()
    ea = r_pr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(r_pr, qn("a:ea"))
    ea.set("typeface", name)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    size: float,
    bold: bool = False,
    color: RGBColor = TEXT_DARK,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    vertical: MSO_ANCHOR = MSO_ANCHOR.MIDDLE,
    word_wrap: bool = True,
    margin: float = 0.04,
    name: str | None = None,
):
    """添加一个只含一个 run 的可编辑文本框。"""
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    if name:
        shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = word_wrap
    frame.vertical_anchor = vertical
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.05
    run = paragraph.add_run()
    run.text = text
    set_east_asian_font(run)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_gradient_background(slide) -> None:
    """添加可编辑的浅绿到中绿渐变底图。"""
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        0,
        0,
        Inches(13.333333),
        Inches(7.5),
    )
    shape.name = "浅绿中绿渐变背景"
    shape.line.fill.background()

    sp_pr = shape._element.spPr
    for tag in ("a:solidFill", "a:gradFill", "a:noFill", "a:pattFill", "a:blipFill"):
        for node in list(sp_pr.findall(qn(tag))):
            sp_pr.remove(node)

    grad_fill = etree.Element(qn("a:gradFill"), rotWithShape="1")
    gs_list = etree.SubElement(grad_fill, qn("a:gsLst"))
    for position, color in (("0", GREEN_LIGHT), ("100000", GREEN_MEDIUM)):
        stop = etree.SubElement(gs_list, qn("a:gs"), pos=position)
        etree.SubElement(stop, qn("a:srgbClr"), val=color)
    etree.SubElement(
        grad_fill,
        qn("a:lin"),
        ang="5400000",
        scaled="1",
    )

    geometry = sp_pr.find(qn("a:prstGeom"))
    insert_at = sp_pr.index(geometry) + 1 if geometry is not None else 1
    sp_pr.insert(insert_at, grad_fill)


def add_accent_motif(slide, section_number: int | None = None) -> None:
    """添加少量可编辑的细节，保持清晰而不压过正文。"""
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0.72),
        Inches(0.40),
        Inches(0.62),
        Inches(0.07),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = ORANGE
    line.line.fill.background()
    if section_number is not None:
        add_text(
            slide,
            f"{section_number:02d} / 09",
            11.55,
            0.24,
            1.05,
            0.36,
            size=11,
            bold=True,
            color=TEXT_MUTED,
            align=PP_ALIGN.RIGHT,
            name="章节进度",
        )


def add_cover(slide, title: str) -> None:
    add_gradient_background(slide)
    add_accent_motif(slide)
    add_text(
        slide,
        "2026 · 北区增长策略",
        1.08,
        1.34,
        11.18,
        0.45,
        size=16,
        bold=True,
        color=TEXT_MUTED,
        align=PP_ALIGN.CENTER,
        name="封面导语",
    )
    add_text(
        slide,
        title,
        0.55,
        2.05,
        12.23,
        1.25,
        size=70,
        bold=True,
        color=TEXT_DARK,
        align=PP_ALIGN.CENTER,
        word_wrap=False,
        margin=0,
        name="封面标题",
    )
    add_text(
        slide,
        "市场机会  ·  产品体系  ·  渠道转化  ·  执行保障",
        1.25,
        3.58,
        10.83,
        0.46,
        size=16,
        color=TEXT_MUTED,
        align=PP_ALIGN.CENTER,
        name="封面摘要",
    )


def add_toc(slide, toc_title: str, entries: list[str]) -> None:
    add_gradient_background(slide)
    add_accent_motif(slide)
    add_text(
        slide,
        toc_title,
        0.85,
        0.67,
        2.0,
        0.75,
        size=36,
        bold=True,
        name="目录标题",
    )
    positions = [
        (0.88, 1.82), (4.48, 1.82), (8.08, 1.82),
        (0.88, 3.33), (4.48, 3.33), (8.08, 3.33),
        (0.88, 4.84), (4.48, 4.84), (8.08, 4.84),
    ]
    for index, (entry, (x, y)) in enumerate(zip(entries, positions), 1):
        add_text(
            slide,
            f"{index:02d}",
            x,
            y,
            0.50,
            0.42,
            size=16,
            bold=True,
            color=ORANGE_TEXT,
            name=f"目录编号{index}",
        )
        add_text(
            slide,
            entry,
            x + 0.57,
            y - 0.03,
            2.78,
            0.70,
            size=18,
            bold=True,
            color=TEXT_DARK,
            name=f"目录条目{index}",
        )


def add_section_page(slide, section_number: int, title: str) -> None:
    add_gradient_background(slide)
    add_accent_motif(slide, section_number)
    add_text(
        slide,
        f"PART {section_number:02d}",
        1.00,
        2.08,
        11.33,
        0.42,
        size=14,
        bold=True,
        color=ORANGE_TEXT,
        align=PP_ALIGN.CENTER,
        name="章节英文标识",
    )
    add_text(
        slide,
        title,
        0.80,
        2.66,
        11.73,
        1.02,
        size=40,
        bold=True,
        color=TEXT_DARK,
        align=PP_ALIGN.CENTER,
        word_wrap=False,
        margin=0,
        name="章节标题",
    )


def add_orange_box(slide, y: float, name: str):
    box = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0.72),
        Inches(y),
        Inches(11.90),
        Inches(2.75),
    )
    box.name = name
    box.fill.background()
    box.line.color.rgb = ORANGE
    box.line.width = Pt(2.25)
    return box


def add_body_page(slide, section_number: int, subsections: list[dict]) -> list[str]:
    add_gradient_background(slide)
    add_accent_motif(slide, section_number)
    shape_ids: list[str] = []
    layouts = ((0.65, 0.90, 1.42), (3.83, 4.08, 4.60))
    for index, (subsection, (box_y, title_y, body_y)) in enumerate(
        zip(subsections, layouts), 1
    ):
        add_orange_box(slide, box_y, f"第{index}节橙色外框")
        title = add_text(
            slide,
            subsection["title"],
            1.02,
            title_y,
            11.28,
            0.42,
            size=18,
            bold=True,
            color=TEXT_DARK,
            align=PP_ALIGN.LEFT,
            name=f"二级标题{index}",
        )
        body = add_text(
            slide,
            subsection["body"],
            1.02,
            body_y,
            11.28,
            1.54,
            size=14,
            color=TEXT_DARK,
            align=PP_ALIGN.LEFT,
            vertical=MSO_ANCHOR.TOP,
            name=f"正文{index}",
        )
        shape_ids.extend((str(title.shape_id), str(body.shape_id)))
    return shape_ids


def template_background_blob(template_path: Path) -> bytes:
    """从真实模板中取出占满整页的背景图，并验证内容。"""
    template = Presentation(str(template_path))
    candidates = []
    for slide in template.slides:
        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            area = int(shape.width or 0) * int(shape.height or 0)
            candidates.append((area, shape.image.blob))
    if not candidates:
        raise ValueError("模板中没有可用的背景图片")
    blob = max(candidates, key=lambda item: item[0])[1]
    if hashlib.sha256(blob).hexdigest() != TEMPLATE_BACKGROUND_SHA256:
        raise ValueError("模板背景图片与逐题检查时的真实输入不一致")
    return blob


def set_transparency(shape, alpha: int) -> None:
    """为纯色填充设置透明度，alpha 取值 0–100000。"""
    solid_fill = shape._element.spPr.find(qn("a:solidFill"))
    if solid_fill is None:
        return
    color = solid_fill.find(qn("a:srgbClr"))
    if color is None:
        return
    etree.SubElement(color, qn("a:alpha"), val=str(alpha))


def add_closing(slide, template_path: Path, closing: dict) -> None:
    background = template_background_blob(template_path)
    picture = slide.shapes.add_picture(
        BytesIO(background),
        0,
        0,
        width=Inches(13.333333),
        height=Inches(7.5),
    )
    picture.name = "来自模版的尾页背景"
    overlay = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        0,
        0,
        Inches(13.333333),
        Inches(7.5),
    )
    overlay.name = "尾页可编辑蒙版"
    overlay.fill.solid()
    overlay.fill.fore_color.rgb = RGBColor(0x0D, 0x29, 0x25)
    set_transparency(overlay, 65000)
    overlay.line.fill.background()
    add_text(
        slide,
        closing["thanks"],
        0.95,
        2.45,
        7.4,
        0.90,
        size=42,
        bold=True,
        color=RGBColor(0xFF, 0xFF, 0xFF),
        name="尾页致谢",
    )
    add_text(
        slide,
        closing["subtitle"],
        0.98,
        3.42,
        7.8,
        0.55,
        size=19,
        color=RGBColor(0xF5, 0xD4, 0x9A),
        name="尾页方案名",
    )


def animation_behavior(parent, timing_id: int, shape_id: str, effect: str) -> int:
    """在当前动画节点下写入显示与入场效果。"""
    child_list = etree.SubElement(parent, f"{{{P_NS}}}childTnLst")
    set_node = etree.SubElement(child_list, f"{{{P_NS}}}set")
    behavior = etree.SubElement(set_node, f"{{{P_NS}}}cBhvr")
    set_ctn = etree.SubElement(
        behavior,
        f"{{{P_NS}}}cTn",
        id=str(timing_id),
        dur="1",
        fill="hold",
    )
    timing_id += 1
    set_conditions = etree.SubElement(set_ctn, f"{{{P_NS}}}stCondLst")
    etree.SubElement(set_conditions, f"{{{P_NS}}}cond", delay="0")
    target = etree.SubElement(behavior, f"{{{P_NS}}}tgtEl")
    etree.SubElement(target, f"{{{P_NS}}}spTgt", spid=shape_id)
    attributes = etree.SubElement(behavior, f"{{{P_NS}}}attrNameLst")
    attribute = etree.SubElement(attributes, f"{{{P_NS}}}attrName")
    attribute.text = "style.visibility"
    destination = etree.SubElement(set_node, f"{{{P_NS}}}to")
    etree.SubElement(destination, f"{{{P_NS}}}strVal", val="visible")

    effect_node = etree.SubElement(
        child_list,
        f"{{{P_NS}}}animEffect",
        transition="in",
        filter=effect,
    )
    effect_behavior = etree.SubElement(effect_node, f"{{{P_NS}}}cBhvr")
    etree.SubElement(
        effect_behavior,
        f"{{{P_NS}}}cTn",
        id=str(timing_id),
        dur="650",
    )
    timing_id += 1
    effect_target = etree.SubElement(effect_behavior, f"{{{P_NS}}}tgtEl")
    etree.SubElement(effect_target, f"{{{P_NS}}}spTgt", spid=shape_id)
    return timing_id


def build_timing(shape_ids: list[str]) -> etree._Element:
    """构造一次单击后按标题、正文、标题、正文自动播放的主序列。"""
    if len(shape_ids) != 4:
        raise ValueError("每张正文页必须有 4 个动画目标")
    timing = etree.Element(f"{{{P_NS}}}timing", nsmap=NSMAP)
    time_nodes = etree.SubElement(timing, f"{{{P_NS}}}tnLst")
    root_parallel = etree.SubElement(time_nodes, f"{{{P_NS}}}par")
    root_ctn = etree.SubElement(
        root_parallel,
        f"{{{P_NS}}}cTn",
        id="1",
        dur="indefinite",
        restart="never",
        nodeType="tmRoot",
    )
    root_children = etree.SubElement(root_ctn, f"{{{P_NS}}}childTnLst")
    sequence = etree.SubElement(
        root_children,
        f"{{{P_NS}}}seq",
        concurrent="1",
        nextAc="seek",
    )
    main_ctn = etree.SubElement(
        sequence,
        f"{{{P_NS}}}cTn",
        id="2",
        dur="indefinite",
        nodeType="mainSeq",
    )
    main_children = etree.SubElement(main_ctn, f"{{{P_NS}}}childTnLst")

    timing_id = 3
    effects = ("fade", "wipe(left)", "fade", "wipe(left)")
    for index, (shape_id, effect) in enumerate(zip(shape_ids, effects)):
        top_parallel = etree.SubElement(main_children, f"{{{P_NS}}}par")
        top_ctn = etree.SubElement(
            top_parallel,
            f"{{{P_NS}}}cTn",
            id=str(timing_id),
            fill="hold",
        )
        timing_id += 1
        top_conditions = etree.SubElement(top_ctn, f"{{{P_NS}}}stCondLst")
        etree.SubElement(
            top_conditions,
            f"{{{P_NS}}}cond",
            delay="indefinite" if index == 0 else "0",
        )
        if index == 0:
            begin = etree.SubElement(
                top_conditions,
                f"{{{P_NS}}}cond",
                evt="onBegin",
                delay="0",
            )
            etree.SubElement(begin, f"{{{P_NS}}}tn", val="2")

        top_children = etree.SubElement(top_ctn, f"{{{P_NS}}}childTnLst")
        behavior_parallel = etree.SubElement(top_children, f"{{{P_NS}}}par")
        behavior_ctn = etree.SubElement(
            behavior_parallel,
            f"{{{P_NS}}}cTn",
            id=str(timing_id),
            fill="hold",
        )
        timing_id += 1
        behavior_conditions = etree.SubElement(
            behavior_ctn,
            f"{{{P_NS}}}stCondLst",
        )
        etree.SubElement(behavior_conditions, f"{{{P_NS}}}cond", delay="0")
        behavior_children = etree.SubElement(
            behavior_ctn,
            f"{{{P_NS}}}childTnLst",
        )
        effect_parallel = etree.SubElement(behavior_children, f"{{{P_NS}}}par")
        effect_ctn = etree.SubElement(
            effect_parallel,
            f"{{{P_NS}}}cTn",
            id=str(timing_id),
            presetID="10" if effect == "fade" else "22",
            presetClass="entr",
            presetSubtype="0",
            fill="hold",
            nodeType="clickEffect" if index == 0 else "afterEffect",
        )
        timing_id += 1
        effect_conditions = etree.SubElement(
            effect_ctn,
            f"{{{P_NS}}}stCondLst",
        )
        etree.SubElement(effect_conditions, f"{{{P_NS}}}cond", delay="0")
        timing_id = animation_behavior(effect_ctn, timing_id, shape_id, effect)

    previous = etree.SubElement(sequence, f"{{{P_NS}}}prevCondLst")
    previous_condition = etree.SubElement(
        previous,
        f"{{{P_NS}}}cond",
        evt="onPrev",
        delay="0",
    )
    previous_target = etree.SubElement(previous_condition, f"{{{P_NS}}}tgtEl")
    etree.SubElement(previous_target, f"{{{P_NS}}}sldTgt")
    following = etree.SubElement(sequence, f"{{{P_NS}}}nextCondLst")
    following_condition = etree.SubElement(
        following,
        f"{{{P_NS}}}cond",
        evt="onNext",
        delay="0",
    )
    following_target = etree.SubElement(following_condition, f"{{{P_NS}}}tgtEl")
    etree.SubElement(following_target, f"{{{P_NS}}}sldTgt")
    return timing


def patch_animations(pptx_path: Path, slide_targets: dict[int, list[str]]) -> None:
    """将动画主序列写入每张正文页的 OOXML。"""
    with tempfile.NamedTemporaryFile(
        suffix=".pptx",
        dir=pptx_path.parent,
        delete=False,
    ) as handle:
        temporary_path = Path(handle.name)
    try:
        with ZipFile(pptx_path, "r") as source, ZipFile(
            temporary_path, "w", ZIP_DEFLATED
        ) as destination:
            for info in source.infolist():
                data = source.read(info.filename)
                slide_no = None
                if info.filename.startswith("ppt/slides/slide") and info.filename.endswith(".xml"):
                    number_text = info.filename.removeprefix("ppt/slides/slide").removesuffix(".xml")
                    if number_text.isdigit():
                        slide_no = int(number_text)
                if slide_no in slide_targets:
                    root = etree.fromstring(data)
                    old_timing = root.find(f"{{{P_NS}}}timing")
                    if old_timing is not None:
                        root.remove(old_timing)
                    timing = build_timing(slide_targets[slide_no])
                    extension_list = root.find(f"{{{P_NS}}}extLst")
                    if extension_list is None:
                        root.append(timing)
                    else:
                        root.insert(root.index(extension_list), timing)
                    data = etree.tostring(
                        root,
                        xml_declaration=True,
                        encoding="UTF-8",
                        standalone=True,
                    )
                destination.writestr(info, data)
        temporary_path.replace(pptx_path)
    finally:
        temporary_path.unlink(missing_ok=True)


def build_deck(template_path: Path, outline_path: Path, output_path: Path) -> None:
    if sha256(template_path) != TEMPLATE_SHA256:
        raise ValueError("输入模板 SHA256 与上游真实文件不一致")
    outline = json.loads(outline_path.read_text(encoding="utf-8"))
    sections = outline["sections"]
    if len(sections) != 9 or any(len(item["subsections"]) != 2 for item in sections):
        raise ValueError("大纲必须严格包含 9 个一级板块，每个板块 2 个二级小节")

    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    blank_layout = presentation.slide_layouts[6]
    while presentation.slides:
        slide_id = presentation.slides._sldIdLst[0]
        presentation.part.drop_rel(slide_id.rId)
        del presentation.slides._sldIdLst[0]

    cover = presentation.slides.add_slide(blank_layout)
    add_cover(cover, outline["title"])
    toc = presentation.slides.add_slide(blank_layout)
    add_toc(toc, outline["toc_title"], [item["toc_title"] for item in sections])

    slide_targets: dict[int, list[str]] = {}
    for section_number, section in enumerate(sections, 1):
        section_slide = presentation.slides.add_slide(blank_layout)
        add_section_page(section_slide, section_number, section["section_title"])
        body_slide = presentation.slides.add_slide(blank_layout)
        slide_no = len(presentation.slides)
        slide_targets[slide_no] = add_body_page(
            body_slide,
            section_number,
            section["subsections"],
        )

    closing_slide = presentation.slides.add_slide(blank_layout)
    add_closing(closing_slide, template_path, outline["closing"])
    if len(presentation.slides) != 21:
        raise RuntimeError(f"交付物应为 21 页，实际为 {len(presentation.slides)} 页")

    presentation.core_properties.title = outline["title"]
    presentation.core_properties.subject = "OmegaUse-OfficeVal officeval_076 逐题参考交付"
    presentation.core_properties.author = "Cowork Evolution Benchmark"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)
    patch_animations(output_path, slide_targets)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument(
        "--outline",
        type=Path,
        default=Path(__file__).with_name("outline.json"),
    )
    arguments = parser.parse_args()
    build_deck(
        arguments.input_dir / TEMPLATE_NAME,
        arguments.outline,
        arguments.output_dir / OUTPUT_NAME,
    )


if __name__ == "__main__":
    main()
