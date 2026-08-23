#!/usr/bin/env python3

SPEC = {'task_id': 'cowork-evo/ppt-005-renewal-review-revision',
 'title': '客户续约评审内容修订',
 'family': 'content_editing',
 'author': '柏舟云服务（合成）',
 'accent': '2F80ED',
 'footer': '柏舟云服务｜内部资料',
 'input_names': ['客户续约评审_待修订.pptx', '续约评审修订清单.md'],
 'source_markers': ['Q2', '84%', '1,380万元', '王芳', '柏舟云服务｜内部资料'],
 'deliverables': [{'name': 'Q3客户续约评审_已修订.pptx', 'type': 'pptx'}],
 'slides': [{'kind': 'cover',
             'title': 'Q3客户续约评审',
             'subtitle': '商务运营会议',
             'bullets': [],
             'kpis': [],
             'chart': None,
             'table': None,
             'timeline': [],
             'callout': '已修订',
             'note': ''},
            {'kind': 'kpi',
             'title': 'Q3续约盘点：续约率提升6个百分点',
             'subtitle': '',
             'bullets': [],
             'kpis': [['84%', '续约率'], ['1,380万元', 'ARR'], ['3家', '高风险客户']],
             'chart': None,
             'table': None,
             'timeline': [],
             'callout': '',
             'note': ''},
            {'kind': 'table',
             'title': '两家重点客户已明确责任人和截止日',
             'subtitle': '',
             'bullets': [],
             'kpis': [],
             'chart': None,
             'table': [['客户', '责任人', '截止日'], ['明辉制造', '王芳', '9月12日'], ['山海药业', '赵晨', '9月15日']],
             'timeline': [],
             'callout': '',
             'note': ''},
            {'kind': 'bullets',
             'title': '续约策略从统一折扣改为分层应对',
             'subtitle': '',
             'bullets': ['绿色：按标准节奏推进续约', '黄色：提前30天复盘价值证据', '红色：责任人与管理层联合恢复'],
             'kpis': [],
             'chart': None,
             'table': None,
             'timeline': [],
             'callout': '',
             'note': ''},
            {'kind': 'closing',
             'title': '下一步：用截止日管理两个重点续约',
             'subtitle': '',
             'bullets': ['王芳：9月12日前完成明辉制造价值复盘', '赵晨：9月15日前完成山海药业方案确认', '周会按续约概率与ARR风险双维跟踪'],
             'kpis': [],
             'chart': None,
             'table': None,
             'timeline': [],
             'callout': '柏舟云服务｜内部资料',
             'note': ''}],
 'content_tokens': ['Q3客户续约评审', '明辉制造', '山海药业', '王芳', '赵晨', '分层应对'],
 'numeric_tokens': ['84%', '1,380万元', '3家', '9月12日', '9月15日', '6个百分点'],
 'preserve_tokens': ['柏舟云服务｜内部资料', '商务运营会议'],
 'forbidden_tokens': ['Q2客户续约评审', '78%', '1,260万元', '陈熙', '李博'],
 'object_mins': {'charts': 0, 'tables': 1, 'pictures': 0, 'shapes': 38},
 'package_tokens': []}

import argparse
from pathlib import Path

import fitz
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


def rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def add_text(slide, text, x, y, w, h, *, size=18, color="243447", bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.06)
    frame.margin_right = Inches(0.06)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.03)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = str(text)
    run.font.name = "Noto Sans CJK SC"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return shape


def add_rect(slide, x, y, w, h, fill, *, radius=True, line=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    return shape


def style_chart(chart, accent):
    chart.has_legend = len(chart.series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.tick_labels.font.size = Pt(10)
    chart.category_axis.tick_labels.font.size = Pt(11)
    colors = [accent, "62B6CB", "F4A261", "7F8C8D"]
    for index, series in enumerate(chart.series):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = rgb(colors[index % len(colors)])


def read_source_text(input_dir: Path) -> str:
    parts = []
    for path in sorted(input_dir.iterdir()):
        suffix = path.suffix.lower()
        if suffix in {".md", ".txt", ".json", ".csv"}:
            parts.append(path.read_text(encoding="utf-8-sig"))
        elif suffix == ".docx":
            document = Document(path)
            parts.extend(paragraph.text for paragraph in document.paragraphs)
            for table in document.tables:
                parts.extend(cell.text for row in table.rows for cell in row.cells)
        elif suffix == ".xlsx":
            workbook = load_workbook(path, data_only=False, read_only=True)
            for sheet in workbook.worksheets:
                parts.extend(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None)
            workbook.close()
        elif suffix == ".pptx":
            presentation = Presentation(path)
            for slide in presentation.slides:
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False):
                        parts.append(shape.text)
                    if getattr(shape, "has_table", False):
                        parts.extend(cell.text for row in shape.table.rows for cell in row.cells)
        elif suffix == ".png":
            parts.append(path.name)
    return "\n".join(parts)


def verify_inputs(input_dir: Path) -> None:
    missing = [name for name in SPEC["input_names"] if not (input_dir / name).is_file()]
    if missing:
        raise FileNotFoundError(f"缺少输入文件：{missing}")
    source_text = read_source_text(input_dir)
    absent = [marker for marker in SPEC["source_markers"] if marker not in source_text]
    if absent:
        raise ValueError(f"输入材料与任务口径不一致，缺少：{absent}")


def build_presentation(output_path: Path, input_dir: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = SPEC["title"]
    prs.core_properties.subject = SPEC["family"]
    prs.core_properties.author = SPEC["author"]
    prs.core_properties.keywords = "CoworkEvoBench, synthetic, PowerPoint"
    accent = SPEC["accent"]
    hero_image = next(iter(sorted(input_dir.glob("*.png"))), None)

    for slide_index, item in enumerate(SPEC["slides"]):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        kind = item["kind"]
        dark = kind in {"cover", "closing"}
        bg = "132238" if dark else ("F3F6F9" if slide_index % 2 else "FFFFFF")
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(bg)

        if kind == "cover":
            add_rect(slide, 0, 0, 0.22, 7.5, accent, radius=False)
            add_rect(slide, 0.8, 0.78, 1.25, 0.13, accent, radius=False)
            if hero_image:
                slide.shapes.add_picture(str(hero_image), Inches(8.7), Inches(0.75), width=Inches(3.9), height=Inches(4.9))
                add_rect(slide, 8.7, 5.75, 3.9, 0.12, accent, radius=False)
                title_width = 7.2
            else:
                for n in range(5):
                    add_rect(slide, 9.1 + n * 0.32, 1.2 + n * 0.55, 2.5 - n * 0.18, 0.18, accent)
                title_width = 8.0
            add_text(slide, item["title"], 0.82, 1.18, title_width, 2.2, size=30, color="FFFFFF", bold=True)
            if item.get("subtitle"):
                add_text(slide, item["subtitle"], 0.82, 3.55, 7.3, 0.75, size=17, color="CBD5E1")
            if item.get("callout"):
                add_rect(slide, 0.82, 5.2, 3.2, 0.58, accent)
                add_text(slide, item["callout"], 0.96, 5.28, 2.92, 0.4, size=14, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER)
        else:
            add_rect(slide, 0, 0, 13.333, 0.12, accent, radius=False)
            add_text(slide, item["title"], 0.72, 0.38, 11.9, 0.82, size=25, color="FFFFFF" if dark else "132238", bold=True)
            add_rect(slide, 0.72, 1.25, 1.05, 0.09, accent, radius=False)

            if kind == "kpi":
                kpis = item.get("kpis", [])
                card_width = 3.72 if len(kpis) >= 3 else 5.7
                for card_index, (value, label) in enumerate(kpis):
                    x = 0.72 + card_index * (card_width + 0.35)
                    add_rect(slide, x, 1.72, card_width, 2.05, "FFFFFF" if not dark else "1D3557", line="DCE4EA")
                    add_text(slide, value, x + 0.18, 2.02, card_width - 0.36, 0.75, size=27, color=accent, bold=True, align=PP_ALIGN.CENTER)
                    add_text(slide, label, x + 0.18, 2.86, card_width - 0.36, 0.45, size=14, color="5E6C84" if not dark else "D7E3F1", align=PP_ALIGN.CENTER)
                for bullet_index, bullet in enumerate(item.get("bullets", [])):
                    add_rect(slide, 0.76, 4.4 + bullet_index * 0.62, 0.18, 0.18, accent)
                    add_text(slide, bullet, 1.05, 4.22 + bullet_index * 0.62, 11.4, 0.55, size=16, color="FFFFFF" if dark else "243447")

            elif kind == "chart":
                chart_spec = item["chart"]
                data = ChartData()
                data.categories = chart_spec["categories"]
                for series_name, values in chart_spec["series"]:
                    data.add_series(series_name, values)
                chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.78), Inches(1.68), Inches(7.4), Inches(4.65), data).chart
                style_chart(chart, accent)
                add_rect(slide, 8.55, 1.68, 4.02, 4.65, "FFFFFF", line="DCE4EA")
                bullets = item.get("bullets") or ["图表口径见输入材料"]
                for bullet_index, bullet in enumerate(bullets):
                    add_rect(slide, 8.9, 2.05 + bullet_index * 1.35, 0.22, 0.22, accent)
                    add_text(slide, bullet, 9.25, 1.82 + bullet_index * 1.35, 2.88, 1.08, size=15)

            elif kind == "table":
                rows = item["table"]
                row_count = len(rows)
                col_count = max(len(row) for row in rows)
                shape = slide.shapes.add_table(row_count, col_count, Inches(0.78), Inches(1.72), Inches(11.78), Inches(3.95))
                table = shape.table
                for row_index, row in enumerate(rows):
                    for col_index in range(col_count):
                        cell = table.cell(row_index, col_index)
                        cell.text = str(row[col_index]) if col_index < len(row) else ""
                        cell.margin_left = Inches(0.08)
                        cell.margin_right = Inches(0.08)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = rgb(accent if row_index == 0 else ("FFFFFF" if row_index % 2 else "EEF3F7"))
                        for paragraph in cell.text_frame.paragraphs:
                            paragraph.alignment = PP_ALIGN.CENTER if row_index == 0 else PP_ALIGN.LEFT
                            for run in paragraph.runs:
                                run.font.name = "Noto Sans CJK SC"
                                run.font.size = Pt(12 if row_index else 13)
                                run.font.bold = row_index == 0
                                run.font.color.rgb = rgb("FFFFFF" if row_index == 0 else "243447")
                if item.get("callout"):
                    add_rect(slide, 0.78, 5.95, 11.78, 0.62, "E8F2F3", line=accent)
                    add_text(slide, item["callout"], 1.0, 6.04, 11.3, 0.4, size=14, bold=True, align=PP_ALIGN.CENTER)

            elif kind == "timeline":
                items = item.get("timeline", [])
                if items:
                    line_y = 3.75
                    add_rect(slide, 1.15, line_y, 10.95, 0.08, accent, radius=False)
                    segment = 10.1 / max(1, len(items) - 1)
                    for item_index, (label, body) in enumerate(items):
                        x = 1.15 + item_index * segment if len(items) > 1 else 5.6
                        add_rect(slide, x, line_y - 0.16, 0.4, 0.4, accent)
                        add_text(slide, label, max(0.45, x - 0.78), 2.0, 1.95, 0.72, size=16, color=accent, bold=True, align=PP_ALIGN.CENTER)
                        add_text(slide, body, max(0.35, x - 1.0), 4.3, 2.35, 1.15, size=14, color="FFFFFF" if dark else "243447", align=PP_ALIGN.CENTER)
                if item.get("callout"):
                    add_rect(slide, 2.8, 6.0, 7.7, 0.55, accent)
                    add_text(slide, item["callout"], 3.0, 6.07, 7.3, 0.36, size=13, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER)

            elif kind in {"bullets", "closing"}:
                for bullet_index, bullet in enumerate(item.get("bullets", [])):
                    y = 1.72 + bullet_index * 1.12
                    add_rect(slide, 0.82, y + 0.08, 0.55, 0.55, accent)
                    add_text(slide, f"{bullet_index + 1:02d}", 0.86, y + 0.13, 0.47, 0.35, size=12, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER)
                    add_text(slide, bullet, 1.62, y - 0.03, 10.4, 0.8, size=18, color="FFFFFF" if dark else "243447", bold=kind == "closing")
                if item.get("callout"):
                    add_rect(slide, 2.2, 6.15, 8.9, 0.58, accent)
                    add_text(slide, item["callout"], 2.4, 6.23, 8.5, 0.38, size=14, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER)

        add_text(slide, SPEC["footer"], 0.72, 7.08, 8.7, 0.24, size=8, color="A9B7C6" if dark else "7F8C8D")
        add_text(slide, f"{slide_index + 1:02d}", 11.95, 7.05, 0.6, 0.25, size=9, color=accent, bold=True, align=PP_ALIGN.RIGHT)
        if item.get("note"):
            try:
                slide.notes_slide.notes_text_frame.text = item["note"]
            except Exception:
                pass

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)


def build_package_files(output_dir: Path) -> None:
    package_tokens = SPEC.get("package_tokens", [])
    if not package_tokens:
        return
    docx_name = next(item["name"] for item in SPEC["deliverables"] if item["type"] == "docx")
    document = Document()
    document.add_heading("2026微光志愿者大会｜讲者手册", 0)
    for heading, body in [
        ("大会主题", "让每一次善意都有回声"),
        ("时间地点", "2026年12月5日，湖滨市民中心二层大礼堂；09:00入场，09:30开始。"),
        ("讲者提示", "年度回顾用数据说明影响；故事环节不公开个人联系方式；明年计划强调志愿者保护。"),
    ]:
        document.add_heading(heading, level=1)
        document.add_paragraph(body)
    document.core_properties.title = "2026微光志愿者大会讲者手册"
    document.core_properties.author = "CoworkEvoBench synthetic solution"
    document.save(output_dir / docx_name)

    pdf_name = next(item["name"] for item in SPEC["deliverables"] if item["type"] == "pdf")
    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)
    lines = [
        "2026微光志愿者大会 参会手册",
        "主题：让每一次善意都有回声",
        "日期：2026年12月5日",
        "地点：湖滨市民中心二层大礼堂",
        "入场：09:00    开始：09:30",
        "无障碍入口：东门",
        "会务提示：请保存本手册，紧急情况联系会务台。",
    ]
    y = 90
    for index, line in enumerate(lines):
        page.insert_text((60, y), line, fontsize=20 if index == 0 else 13, fontname="china-s", color=(0.08, 0.13, 0.22))
        y += 62 if index == 0 else 38
    page.draw_rect(fitz.Rect(48, 48, 547, 794), color=(0.9, 0.49, 0.13), width=2)
    pdf.set_metadata({"title": "2026微光志愿者大会参会手册", "author": "CoworkEvoBench synthetic solution"})
    pdf.save(output_dir / pdf_name)
    pdf.close()


def main() -> int:
    parser = argparse.ArgumentParser(description="生成本题的标准交付物")
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    verify_inputs(args.input_dir)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    pptx_name = next(item["name"] for item in SPEC["deliverables"] if item["type"] == "pptx")
    build_presentation(args.output_dir / pptx_name, args.input_dir)
    build_package_files(args.output_dir)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
