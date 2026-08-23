#!/usr/bin/env python3

SPEC = {'task_id': 'cowork-evo/ppt-005-renewal-review-revision',
 'title': '客户续约评审内容修订',
 'family': 'content_editing',
 'author': '柏舟云服务（合成）',
 'accent': '2F80ED',
 'footer': '柏舟云服务｜内部资料',
 'input_names': ['客户续约评审_待修订.pptx', '续约评审修订清单.md'],
 'source_markers': ['Q2', '84%', '1,380万元', '王芳', '柏舟云服务｜内部资料'],
 'deliverables': [{'name': 'Q3客户续约评审_已修订.pptx', 'type': 'pptx'}],
 'slides': [{'kind': 'cover',
             'title': 'Q3客户续约评审',
             'subtitle': '商务运营会议',
             'bullets': [],
             'kpis': [],
             'chart': None,
             'table': None,
             'timeline': [],
             'callout': '已修订',
             'note': ''},
            {'kind': 'kpi',
             'title': 'Q3续约盘点：续约率提升6个百分点',
             'subtitle': '',
             'bullets': [],
             'kpis': [['84%', '续约率'], ['1,380万元', 'ARR'], ['3家', '高风险客户']],
             'chart': None,
             'table': None,
             'timeline': [],
             'callout': '',
             'note': ''},
            {'kind': 'table',
             'title': '两家重点客户已明确责任人和截止日',
             'subtitle': '',
             'bullets': [],
             'kpis': [],
             'chart': None,
             'table': [['客户', '责任人', '截止日'], ['明辉制造', '王芳', '9月12日'], ['山海药业', '赵晨', '9月15日']],
             'timeline': [],
             'callout': '',
             'note': ''},
            {'kind': 'bullets',
             'title': '续约策略从统一折扣改为分层应对',
             'subtitle': '',
             'bullets': ['绿色：按标准节奏推进续约', '黄色：提前30天复盘价值证据', '红色：责任人与管理层联合恢复'],
             'kpis': [],
             'chart': None,
             'table': None,
             'timeline': [],
             'callout': '',
             'note': ''},
            {'kind': 'closing',
             'title': '下一步：用截止日管理两个重点续约',
             'subtitle': '',
             'bullets': ['王芳：9月12日前完成明辉制造价值复盘', '赵晨：9月15日前完成山海药业方案确认', '周会按续约概率与ARR风险双维跟踪'],
             'kpis': [],
             'chart': None,
             'table': None,
             'timeline': [],
             'callout': '柏舟云服务｜内部资料',
             'note': ''}],
 'content_tokens': ['Q3客户续约评审', '明辉制造', '山海药业', '王芳', '赵晨', '分层应对'],
 'numeric_tokens': ['84%', '1,380万元', '3家', '9月12日', '9月15日', '6个百分点'],
 'preserve_tokens': ['柏舟云服务｜内部资料', '商务运营会议'],
 'forbidden_tokens': ['Q2客户续约评审', '78%', '1,260万元', '陈熙', '李博'],
 'object_mins': {'charts': 0, 'tables': 1, 'pictures': 0, 'shapes': 38},
 'package_tokens': []}

import argparse
import hashlib
import json
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


WEIGHTS = {"R001": 3, "R002": 2, "R003": 4, "R004": 3, "R005": 2, "R006": 2, "R007": -3, "R008": 1}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def ratio(tokens, text):
    if not tokens:
        return 1.0, []
    missing = [token for token in tokens if token not in text]
    return (len(tokens) - len(missing)) / len(tokens), missing


def inspect_presentation(path: Path):
    prs = Presentation(path)
    slide_texts = []
    charts = tables = pictures = shapes = accent_shapes = out_of_bounds = 0
    font_sizes = []
    for slide in prs.slides:
        texts = []
        for shape in slide.shapes:
            shapes += 1
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width or shape.top + shape.height > prs.slide_height:
                out_of_bounds += 1
            if getattr(shape, "has_chart", False):
                charts += 1
            if getattr(shape, "has_table", False):
                tables += 1
                texts.extend(cell.text for row in shape.table.rows for cell in row.cells)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
            if getattr(shape, "has_text_frame", False):
                texts.append(shape.text)
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.text.strip() and run.font.size:
                            font_sizes.append(run.font.size.pt)
            try:
                if shape.fill.type and str(shape.fill.fore_color.rgb).upper() == SPEC["accent"].upper():
                    accent_shapes += 1
            except Exception:
                pass
        slide_texts.append("\n".join(texts))
    return {
        "presentation": prs,
        "slide_texts": slide_texts,
        "text": "\n".join(slide_texts),
        "charts": charts,
        "tables": tables,
        "pictures": pictures,
        "shapes": shapes,
        "accent_shapes": accent_shapes,
        "out_of_bounds": out_of_bounds,
        "min_font": min(font_sizes) if font_sizes else 0,
    }


def criterion(criterion_id, score, evidence):
    return {"criterion_id": criterion_id, "score": round(max(0.0, min(1.0, score)), 4), "evidence": evidence, "method": "deterministic-pptx-v1"}


def evaluate(artifact_dir: Path):
    deliverable_paths = [artifact_dir / item["name"] for item in SPEC["deliverables"]]
    present_paths = [path for path in deliverable_paths if path.is_file()]
    hashes = [sha256(path) for path in present_paths]
    if not hashes:
        hashes = [hashlib.sha256(f"missing:{SPEC['task_id']}".encode()).hexdigest()]

    missing_files = [path.name for path in deliverable_paths if not path.is_file()]
    ppt_item = next(item for item in SPEC["deliverables"] if item["type"] == "pptx")
    ppt_path = artifact_dir / ppt_item["name"]
    inspection = None
    open_errors = []
    if ppt_path.is_file():
        try:
            inspection = inspect_presentation(ppt_path)
        except Exception as error:
            open_errors.append(f"PPTX无法打开: {error}")

    for item in SPEC["deliverables"]:
        path = artifact_dir / item["name"]
        if not path.is_file():
            continue
        try:
            if item["type"] == "docx":
                document = Document(path)
                doc_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
                absent = [token for token in SPEC.get("package_tokens", []) if token not in doc_text]
                if absent:
                    open_errors.append(f"DOCX缺少必需内容: {absent}")
            elif item["type"] == "pdf":
                if path.stat().st_size < 800 or path.read_bytes()[:5] != b"%PDF-":
                    open_errors.append("PDF文件不完整")
        except Exception as error:
            open_errors.append(f"{item['type'].upper()}无法打开: {error}")

    file_score = 1.0 if not missing_files and not open_errors and inspection else 0.0
    criteria = [criterion("R001", file_score, f"缺失={missing_files or '无'}；打开错误={open_errors or '无'}")]
    if inspection is None:
        for criterion_id in ["R002", "R003", "R004", "R005", "R006"]:
            criteria.append(criterion(criterion_id, 0.0, "无可解析的PPTX产物"))
        criteria.append(criterion("R007", 1.0, "缺少有效PPTX视为严重负向违规"))
        criteria.append(criterion("R008", 0.0, "无法校验格式保留项"))
    else:
        slide_count = len(inspection["slide_texts"])
        expected_titles = [item["title"] for item in SPEC["slides"]]
        title_hits = sum(
            index < slide_count and title in inspection["slide_texts"][index]
            for index, title in enumerate(expected_titles)
        )
        title_score = (title_hits / len(expected_titles)) if slide_count == len(expected_titles) else min(0.5, title_hits / len(expected_titles))
        criteria.append(criterion("R002", title_score, f"页数={slide_count}/{len(expected_titles)}；按顺序命中标题={title_hits}/{len(expected_titles)}"))

        content_score, missing_content = ratio(SPEC["content_tokens"], inspection["text"])
        criteria.append(criterion("R003", content_score, f"缺少内容={missing_content or '无'}"))
        data_score, missing_numbers = ratio(SPEC["numeric_tokens"], inspection["text"])
        criteria.append(criterion("R004", data_score, f"缺少数据={missing_numbers or '无'}"))

        mins = SPEC["object_mins"]
        ratios = []
        for key in ["charts", "tables", "pictures", "shapes"]:
            required = mins.get(key, 0)
            ratios.append(1.0 if required == 0 else min(1.0, inspection[key] / required))
        structure_score = sum(ratios) / len(ratios)
        criteria.append(criterion("R005", structure_score, f"图表={inspection['charts']}，表格={inspection['tables']}，图片={inspection['pictures']}，形状={inspection['shapes']}"))

        visual_checks = [inspection["out_of_bounds"] == 0, inspection["min_font"] >= 8, inspection["accent_shapes"] >= len(SPEC["slides"])]
        visual_score = sum(visual_checks) / len(visual_checks)
        criteria.append(criterion("R006", visual_score, f"越界对象={inspection['out_of_bounds']}，最小字号={inspection['min_font']:.1f}pt，强调色对象={inspection['accent_shapes']}"))

        forbidden_hits = [token for token in SPEC["forbidden_tokens"] if token in inspection["text"]]
        blank_slides = [index + 1 for index, text in enumerate(inspection["slide_texts"]) if len(text.strip()) < 12]
        violation_score = min(1.0, (len(forbidden_hits) + len(blank_slides)) / 2)
        criteria.append(criterion("R007", violation_score, f"禁止内容={forbidden_hits or '无'}；空白页={blank_slides or '无'}"))

        preserve_score, missing_preserve = ratio(SPEC["preserve_tokens"], inspection["text"])
        wide = abs(inspection["presentation"].slide_width / inspection["presentation"].slide_height - 16 / 9) < 0.02
        format_score = (preserve_score + float(wide)) / 2
        criteria.append(criterion("R008", format_score, f"缺少保留项={missing_preserve or '无'}；16:9={wide}"))

    positive_weight = sum(weight for weight in WEIGHTS.values() if weight > 0)
    scores = {item["criterion_id"]: item["score"] for item in criteria}
    earned = sum(weight * scores[criterion_id] for criterion_id, weight in WEIGHTS.items() if weight > 0)
    penalty = sum(abs(weight) * scores[criterion_id] for criterion_id, weight in WEIGHTS.items() if weight < 0)
    reward = max(0.0, min(1.0, (earned - penalty) / positive_weight))
    hurdle_failures = [] if scores["R001"] == 1.0 else ["R001"]
    if hurdle_failures:
        reward = 0.0
    result = {
        "task_id": SPEC["task_id"],
        "rubric_version": "cowork-evo-rubric-v0.1",
        "judge_version": "ppt-deterministic-v1.0",
        "artifact_sha256": sorted(set(hashes)),
        "criterion_results": criteria,
        "reward": round(reward, 4),
        "passed": not hurdle_failures and reward >= 0.7,
        "hurdle_failures": hurdle_failures,
    }
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description="对PPT任务产物做确定性评分")
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    args = parser.parse_args()
    result = evaluate(args.artifact_dir)
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"task_id": result["task_id"], "reward": result["reward"], "passed": result["passed"]}, ensure_ascii=False))
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
