#!/usr/bin/env python3

SPEC = {'task_id': 'cowork-evo/ppt-007-quarterly-ops-polish',
 'title': '季度经营会演示稿润色',
 'family': 'deck_polishing',
 'author': '柚子生活（合成）',
 'accent': '0C7C86',
 'footer': '柚子生活｜内部经营会',
 'input_names': ['二季度经营会_初稿.pptx', '经营会润色要求.md'],
 'source_markers': ['2,480万元', '41.2%', '15.6万', '1.8个百分点', '柚子生活｜内部经营会'],
 'deliverables': [{'name': '2026Q2经营会_润色版.pptx', 'type': 'pptx'}],
 'slides': [{'kind': 'cover',
             'title': '2026 Q2：收入增长，但拉新与复购同时承压',
             'subtitle': '柚子生活季度经营会',
             'bullets': [],
             'kpis': [],
             'chart': None,
             'table': None,
             'timeline': [],
             'callout': 'Q2 REVIEW',
             'note': ''},
            {'kind': 'kpi',
             'title': '2,480 万元收入达成，毛利率保持 41.2%',
             'subtitle': '',
             'bullets': [],
             'kpis': [['2,480万元', '收入'], ['41.2%', '贡献毛利率'], ['15.6万', '付费会员']],
             'chart': None,
             'table': None,
             'timeline': [],
             'callout': '',
             'note': ''},
            {'kind': 'chart',
             'title': '6月收入 910 万元，连续三月增长',
             'subtitle': '',
             'bullets': ['增长主要由6月会员日与门店联动带动', '下季度不将单次促销外推为自然增长'],
             'kpis': [],
             'chart': {'categories': ['4月', '5月', '6月'], 'series': [['收入(万元)', [760, 810, 910]]]},
             'table': None,
             'timeline': [],
             'callout': '',
             'note': ''},
            {'kind': 'table',
             'title': '深圳贡献最高，厦门利润率仅 11.9%',
             'subtitle': '',
             'bullets': [],
             'kpis': [],
             'chart': None,
             'table': [['城市', '收入', '利润率', '判断'],
                       ['深圳', '860万元', '13.8%', '保持'],
                       ['广州', '740万元', '12.6%', '提效'],
                       ['厦门', '420万元', '11.9%', '重点改善']],
             'timeline': [],
             'callout': '',
             'note': ''},
            {'kind': 'bullets',
             'title': '三个问题都已找到可量化的管理抓手',
             'subtitle': '',
             'bullets': ['拉新：将付费渠道CAC上限设为68元', '人效：厦门排班按小时客流重算', '复购：针对60天未访会员做分层召回'],
             'kpis': [],
             'chart': None,
             'table': None,
             'timeline': [],
             'callout': '复购率较上季度下降1.8个百分点',
             'note': ''},
            {'kind': 'closing',
             'title': 'Q3：责任到人，指标到周',
             'subtitle': '',
             'bullets': ['市场负责人林晴：将付费渠道CAC降至68元以内', '区域负责人周扬：将厦门利润率提至12.5%', '会员负责人许文：在9月末前恢复1.0个百分点复购率'],
             'kpis': [],
             'chart': None,
             'table': None,
             'timeline': [],
             'callout': '柚子生活｜内部经营会',
             'note': ''}],
 'content_tokens': ['柚子生活季度经营会', '厦门', '分层召回', '责任到人', '指标到周'],
 'numeric_tokens': ['2,480万元', '41.2%', '15.6万', '910 万元', '11.9%', '1.8个百分点', '68元', '12.5%'],
 'preserve_tokens': ['柚子生活｜内部经营会', 'Q2 REVIEW'],
 'forbidden_tokens': ['数据待定', '责任人待定', '保证恢复'],
 'object_mins': {'charts': 1, 'tables': 1, 'pictures': 0, 'shapes': 46},
 'package_tokens': []}

import argparse
import hashlib
import json
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


WEIGHTS = {"R001": 3, "R002": 2, "R003": 4, "R004": 3, "R005": 2, "R006": 2, "R007": -3, "R008": 1}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def ratio(tokens, text):
    if not tokens:
        return 1.0, []
    missing = [token for token in tokens if token not in text]
    return (len(tokens) - len(missing)) / len(tokens), missing


def inspect_presentation(path: Path):
    prs = Presentation(path)
    slide_texts = []
    charts = tables = pictures = shapes = accent_shapes = out_of_bounds = 0
    font_sizes = []
    for slide in prs.slides:
        texts = []
        for shape in slide.shapes:
            shapes += 1
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width or shape.top + shape.height > prs.slide_height:
                out_of_bounds += 1
            if getattr(shape, "has_chart", False):
                charts += 1
            if getattr(shape, "has_table", False):
                tables += 1
                texts.extend(cell.text for row in shape.table.rows for cell in row.cells)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
            if getattr(shape, "has_text_frame", False):
                texts.append(shape.text)
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.text.strip() and run.font.size:
                            font_sizes.append(run.font.size.pt)
            try:
                if shape.fill.type and str(shape.fill.fore_color.rgb).upper() == SPEC["accent"].upper():
                    accent_shapes += 1
            except Exception:
                pass
        slide_texts.append("\n".join(texts))
    return {
        "presentation": prs,
        "slide_texts": slide_texts,
        "text": "\n".join(slide_texts),
        "charts": charts,
        "tables": tables,
        "pictures": pictures,
        "shapes": shapes,
        "accent_shapes": accent_shapes,
        "out_of_bounds": out_of_bounds,
        "min_font": min(font_sizes) if font_sizes else 0,
    }


def criterion(criterion_id, score, evidence):
    return {"criterion_id": criterion_id, "score": round(max(0.0, min(1.0, score)), 4), "evidence": evidence, "method": "deterministic-pptx-v1"}


def evaluate(artifact_dir: Path):
    deliverable_paths = [artifact_dir / item["name"] for item in SPEC["deliverables"]]
    present_paths = [path for path in deliverable_paths if path.is_file()]
    hashes = [sha256(path) for path in present_paths]
    if not hashes:
        hashes = [hashlib.sha256(f"missing:{SPEC['task_id']}".encode()).hexdigest()]

    missing_files = [path.name for path in deliverable_paths if not path.is_file()]
    ppt_item = next(item for item in SPEC["deliverables"] if item["type"] == "pptx")
    ppt_path = artifact_dir / ppt_item["name"]
    inspection = None
    open_errors = []
    if ppt_path.is_file():
        try:
            inspection = inspect_presentation(ppt_path)
        except Exception as error:
            open_errors.append(f"PPTX无法打开: {error}")

    for item in SPEC["deliverables"]:
        path = artifact_dir / item["name"]
        if not path.is_file():
            continue
        try:
            if item["type"] == "docx":
                document = Document(path)
                doc_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
                absent = [token for token in SPEC.get("package_tokens", []) if token not in doc_text]
                if absent:
                    open_errors.append(f"DOCX缺少必需内容: {absent}")
            elif item["type"] == "pdf":
                if path.stat().st_size < 800 or path.read_bytes()[:5] != b"%PDF-":
                    open_errors.append("PDF文件不完整")
        except Exception as error:
            open_errors.append(f"{item['type'].upper()}无法打开: {error}")

    file_score = 1.0 if not missing_files and not open_errors and inspection else 0.0
    criteria = [criterion("R001", file_score, f"缺失={missing_files or '无'}；打开错误={open_errors or '无'}")]
    if inspection is None:
        for criterion_id in ["R002", "R003", "R004", "R005", "R006"]:
            criteria.append(criterion(criterion_id, 0.0, "无可解析的PPTX产物"))
        criteria.append(criterion("R007", 1.0, "缺少有效PPTX视为严重负向违规"))
        criteria.append(criterion("R008", 0.0, "无法校验格式保留项"))
    else:
        slide_count = len(inspection["slide_texts"])
        expected_titles = [item["title"] for item in SPEC["slides"]]
        title_hits = sum(
            index < slide_count and title in inspection["slide_texts"][index]
            for index, title in enumerate(expected_titles)
        )
        title_score = (title_hits / len(expected_titles)) if slide_count == len(expected_titles) else min(0.5, title_hits / len(expected_titles))
        criteria.append(criterion("R002", title_score, f"页数={slide_count}/{len(expected_titles)}；按顺序命中标题={title_hits}/{len(expected_titles)}"))

        content_score, missing_content = ratio(SPEC["content_tokens"], inspection["text"])
        criteria.append(criterion("R003", content_score, f"缺少内容={missing_content or '无'}"))
        data_score, missing_numbers = ratio(SPEC["numeric_tokens"], inspection["text"])
        criteria.append(criterion("R004", data_score, f"缺少数据={missing_numbers or '无'}"))

        mins = SPEC["object_mins"]
        ratios = []
        for key in ["charts", "tables", "pictures", "shapes"]:
            required = mins.get(key, 0)
            ratios.append(1.0 if required == 0 else min(1.0, inspection[key] / required))
        structure_score = sum(ratios) / len(ratios)
        criteria.append(criterion("R005", structure_score, f"图表={inspection['charts']}，表格={inspection['tables']}，图片={inspection['pictures']}，形状={inspection['shapes']}"))

        visual_checks = [inspection["out_of_bounds"] == 0, inspection["min_font"] >= 8, inspection["accent_shapes"] >= len(SPEC["slides"])]
        visual_score = sum(visual_checks) / len(visual_checks)
        criteria.append(criterion("R006", visual_score, f"越界对象={inspection['out_of_bounds']}，最小字号={inspection['min_font']:.1f}pt，强调色对象={inspection['accent_shapes']}"))

        forbidden_hits = [token for token in SPEC["forbidden_tokens"] if token in inspection["text"]]
        blank_slides = [index + 1 for index, text in enumerate(inspection["slide_texts"]) if len(text.strip()) < 12]
        violation_score = min(1.0, (len(forbidden_hits) + len(blank_slides)) / 2)
        criteria.append(criterion("R007", violation_score, f"禁止内容={forbidden_hits or '无'}；空白页={blank_slides or '无'}"))

        preserve_score, missing_preserve = ratio(SPEC["preserve_tokens"], inspection["text"])
        wide = abs(inspection["presentation"].slide_width / inspection["presentation"].slide_height - 16 / 9) < 0.02
        format_score = (preserve_score + float(wide)) / 2
        criteria.append(criterion("R008", format_score, f"缺少保留项={missing_preserve or '无'}；16:9={wide}"))

    positive_weight = sum(weight for weight in WEIGHTS.values() if weight > 0)
    scores = {item["criterion_id"]: item["score"] for item in criteria}
    earned = sum(weight * scores[criterion_id] for criterion_id, weight in WEIGHTS.items() if weight > 0)
    penalty = sum(abs(weight) * scores[criterion_id] for criterion_id, weight in WEIGHTS.items() if weight < 0)
    reward = max(0.0, min(1.0, (earned - penalty) / positive_weight))
    hurdle_failures = [] if scores["R001"] == 1.0 else ["R001"]
    if hurdle_failures:
        reward = 0.0
    result = {
        "task_id": SPEC["task_id"],
        "rubric_version": "cowork-evo-rubric-v0.1",
        "judge_version": "ppt-deterministic-v1.0",
        "artifact_sha256": sorted(set(hashes)),
        "criterion_results": criteria,
        "reward": round(reward, 4),
        "passed": not hurdle_failures and reward >= 0.7,
        "hurdle_failures": hurdle_failures,
    }
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description="对PPT任务产物做确定性评分")
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    args = parser.parse_args()
    result = evaluate(args.artifact_dir)
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"task_id": result["task_id"], "reward": result["reward"], "passed": result["passed"]}, ensure_ascii=False))
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
