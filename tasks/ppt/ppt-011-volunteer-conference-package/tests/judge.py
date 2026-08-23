#!/usr/bin/env python3

SPEC = {'task_id': 'cowork-evo/ppt-011-volunteer-conference-package',
 'title': '志愿者大会演示交付包',
 'family': 'presentation_package',
 'author': '微光公益联盟（合成）',
 'accent': 'E67E22',
 'footer': '微光公益联盟｜让善意被看见',
 'input_names': ['年度志愿者大会素材.docx', '参会数据.xlsx', '会务信息.md', '大会主视觉.png'],
 'source_markers': ['1,280', '32,640', '8,900', '2026年12月5日', '400'],
 'deliverables': [{'name': '2026微光志愿者大会.pptx', 'type': 'pptx'},
                  {'name': '2026微光志愿者大会_参会手册.pdf', 'type': 'pdf'},
                  {'name': '2026微光志愿者大会_讲者手册.docx', 'type': 'docx'}],
 'slides': [{'kind': 'cover',
             'title': '让每一次善意都有回声',
             'subtitle': '2026微光志愿者大会',
             'bullets': [],
             'kpis': [],
             'chart': None,
             'table': None,
             'timeline': [],
             'callout': '2026年12月5日',
             'note': ''},
            {'kind': 'kpi',
             'title': '1,280 名志愿者，用 32,640 小时服务 21 个社区',
             'subtitle': '',
             'bullets': [],
             'kpis': [['1,280名', '活跃志愿者'], ['32,640小时', '志愿服务'], ['8,900人次', '受益者']],
             'chart': None,
             'table': None,
             'timeline': [],
             'callout': '',
             'note': ''},
            {'kind': 'chart',
             'title': '400 位参会者中，志愿者占 71.5%',
             'subtitle': '',
             'bullets': ['设置志愿者快速签到通道', '社区伙伴与捐赠方席位相邻，便于交流'],
             'kpis': [],
             'chart': {'categories': ['志愿者', '社区伙伴', '捐赠方', '工作人员'], 'series': [['人数', [286, 72, 24, 18]]]},
             'table': None,
             'timeline': [],
             'callout': '',
             'note': ''},
            {'kind': 'table',
             'title': '三个核心环节在 11:30 前完成',
             'subtitle': '',
             'bullets': [],
             'kpis': [],
             'chart': None,
             'table': [['时间', '环节', '负责'],
                       ['09:30', '年度回顾', '项目组'],
                       ['10:10', '志愿者故事', '社区伙伴'],
                       ['11:00', '明年计划', '秘书处']],
             'timeline': [],
             'callout': '',
             'note': ''},
            {'kind': 'bullets',
             'title': '夜间陪诊从 80 人扩展至 230 人',
             'subtitle': '',
             'bullets': ['增长来自社区转介与志愿者口碑', '扩展后的首要问题是排班可持续性', '明年将志愿者保护与心理支持纳入服务标准'],
             'kpis': [],
             'chart': None,
             'table': None,
             'timeline': [],
             'callout': '',
             'note': ''},
            {'kind': 'timeline',
             'title': '2027：新增 6 个服务点，将核心留存率提升至 80%',
             'subtitle': '',
             'bullets': [],
             'kpis': [],
             'chart': None,
             'table': None,
             'timeline': [['Q1', '完成需求地图与合作社区签约'], ['Q2', '启动3个服务点'], ['Q3–Q4', '完成其余3个点并复盘留存']],
             'callout': '',
             'note': ''},
            {'kind': 'closing',
             'title': '感谢每一位让善意产生回声的人',
             'subtitle': '',
             'bullets': ['地点：湖滨市民中心二层大礼堂', '入场：09:00，正式开始：09:30', '无障碍入口：东门；紧急情况联系会务台'],
             'kpis': [],
             'chart': None,
             'table': None,
             'timeline': [],
             'callout': '微光公益联盟｜让善意被看见',
             'note': ''}],
 'content_tokens': ['让每一次善意都有回声', '夜间陪诊', '志愿者保护', '湖滨市民中心二层大礼堂', '无障碍入口'],
 'numeric_tokens': ['1,280名', '32,640小时', '8,900人次', '400 ', '71.5%', '80 人', '230 人', '80%'],
 'preserve_tokens': ['2026微光志愿者大会', '2026年12月5日'],
 'forbidden_tokens': ['个人手机号', '真实住址', '身份证号'],
 'object_mins': {'charts': 1, 'tables': 1, 'pictures': 1, 'shapes': 52},
 'package_tokens': ['让每一次善意都有回声', '2026年12月5日', '湖滨市民中心', '09:30']}

import argparse
import hashlib
import json
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


WEIGHTS = {"R001": 3, "R002": 2, "R003": 4, "R004": 3, "R005": 2, "R006": 2, "R007": -3, "R008": 1}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def ratio(tokens, text):
    if not tokens:
        return 1.0, []
    missing = [token for token in tokens if token not in text]
    return (len(tokens) - len(missing)) / len(tokens), missing


def inspect_presentation(path: Path):
    prs = Presentation(path)
    slide_texts = []
    charts = tables = pictures = shapes = accent_shapes = out_of_bounds = 0
    font_sizes = []
    for slide in prs.slides:
        texts = []
        for shape in slide.shapes:
            shapes += 1
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width or shape.top + shape.height > prs.slide_height:
                out_of_bounds += 1
            if getattr(shape, "has_chart", False):
                charts += 1
            if getattr(shape, "has_table", False):
                tables += 1
                texts.extend(cell.text for row in shape.table.rows for cell in row.cells)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
            if getattr(shape, "has_text_frame", False):
                texts.append(shape.text)
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.text.strip() and run.font.size:
                            font_sizes.append(run.font.size.pt)
            try:
                if shape.fill.type and str(shape.fill.fore_color.rgb).upper() == SPEC["accent"].upper():
                    accent_shapes += 1
            except Exception:
                pass
        slide_texts.append("\n".join(texts))
    return {
        "presentation": prs,
        "slide_texts": slide_texts,
        "text": "\n".join(slide_texts),
        "charts": charts,
        "tables": tables,
        "pictures": pictures,
        "shapes": shapes,
        "accent_shapes": accent_shapes,
        "out_of_bounds": out_of_bounds,
        "min_font": min(font_sizes) if font_sizes else 0,
    }


def criterion(criterion_id, score, evidence):
    return {"criterion_id": criterion_id, "score": round(max(0.0, min(1.0, score)), 4), "evidence": evidence, "method": "deterministic-pptx-v1"}


def evaluate(artifact_dir: Path):
    deliverable_paths = [artifact_dir / item["name"] for item in SPEC["deliverables"]]
    present_paths = [path for path in deliverable_paths if path.is_file()]
    hashes = [sha256(path) for path in present_paths]
    if not hashes:
        hashes = [hashlib.sha256(f"missing:{SPEC['task_id']}".encode()).hexdigest()]

    missing_files = [path.name for path in deliverable_paths if not path.is_file()]
    ppt_item = next(item for item in SPEC["deliverables"] if item["type"] == "pptx")
    ppt_path = artifact_dir / ppt_item["name"]
    inspection = None
    open_errors = []
    if ppt_path.is_file():
        try:
            inspection = inspect_presentation(ppt_path)
        except Exception as error:
            open_errors.append(f"PPTX无法打开: {error}")

    for item in SPEC["deliverables"]:
        path = artifact_dir / item["name"]
        if not path.is_file():
            continue
        try:
            if item["type"] == "docx":
                document = Document(path)
                doc_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
                absent = [token for token in SPEC.get("package_tokens", []) if token not in doc_text]
                if absent:
                    open_errors.append(f"DOCX缺少必需内容: {absent}")
            elif item["type"] == "pdf":
                if path.stat().st_size < 800 or path.read_bytes()[:5] != b"%PDF-":
                    open_errors.append("PDF文件不完整")
        except Exception as error:
            open_errors.append(f"{item['type'].upper()}无法打开: {error}")

    file_score = 1.0 if not missing_files and not open_errors and inspection else 0.0
    criteria = [criterion("R001", file_score, f"缺失={missing_files or '无'}；打开错误={open_errors or '无'}")]
    if inspection is None:
        for criterion_id in ["R002", "R003", "R004", "R005", "R006"]:
            criteria.append(criterion(criterion_id, 0.0, "无可解析的PPTX产物"))
        criteria.append(criterion("R007", 1.0, "缺少有效PPTX视为严重负向违规"))
        criteria.append(criterion("R008", 0.0, "无法校验格式保留项"))
    else:
        slide_count = len(inspection["slide_texts"])
        expected_titles = [item["title"] for item in SPEC["slides"]]
        title_hits = sum(
            index < slide_count and title in inspection["slide_texts"][index]
            for index, title in enumerate(expected_titles)
        )
        title_score = (title_hits / len(expected_titles)) if slide_count == len(expected_titles) else min(0.5, title_hits / len(expected_titles))
        criteria.append(criterion("R002", title_score, f"页数={slide_count}/{len(expected_titles)}；按顺序命中标题={title_hits}/{len(expected_titles)}"))

        content_score, missing_content = ratio(SPEC["content_tokens"], inspection["text"])
        criteria.append(criterion("R003", content_score, f"缺少内容={missing_content or '无'}"))
        data_score, missing_numbers = ratio(SPEC["numeric_tokens"], inspection["text"])
        criteria.append(criterion("R004", data_score, f"缺少数据={missing_numbers or '无'}"))

        mins = SPEC["object_mins"]
        ratios = []
        for key in ["charts", "tables", "pictures", "shapes"]:
            required = mins.get(key, 0)
            ratios.append(1.0 if required == 0 else min(1.0, inspection[key] / required))
        structure_score = sum(ratios) / len(ratios)
        criteria.append(criterion("R005", structure_score, f"图表={inspection['charts']}，表格={inspection['tables']}，图片={inspection['pictures']}，形状={inspection['shapes']}"))

        visual_checks = [inspection["out_of_bounds"] == 0, inspection["min_font"] >= 8, inspection["accent_shapes"] >= len(SPEC["slides"])]
        visual_score = sum(visual_checks) / len(visual_checks)
        criteria.append(criterion("R006", visual_score, f"越界对象={inspection['out_of_bounds']}，最小字号={inspection['min_font']:.1f}pt，强调色对象={inspection['accent_shapes']}"))

        forbidden_hits = [token for token in SPEC["forbidden_tokens"] if token in inspection["text"]]
        blank_slides = [index + 1 for index, text in enumerate(inspection["slide_texts"]) if len(text.strip()) < 12]
        violation_score = min(1.0, (len(forbidden_hits) + len(blank_slides)) / 2)
        criteria.append(criterion("R007", violation_score, f"禁止内容={forbidden_hits or '无'}；空白页={blank_slides or '无'}"))

        preserve_score, missing_preserve = ratio(SPEC["preserve_tokens"], inspection["text"])
        wide = abs(inspection["presentation"].slide_width / inspection["presentation"].slide_height - 16 / 9) < 0.02
        format_score = (preserve_score + float(wide)) / 2
        criteria.append(criterion("R008", format_score, f"缺少保留项={missing_preserve or '无'}；16:9={wide}"))

    positive_weight = sum(weight for weight in WEIGHTS.values() if weight > 0)
    scores = {item["criterion_id"]: item["score"] for item in criteria}
    earned = sum(weight * scores[criterion_id] for criterion_id, weight in WEIGHTS.items() if weight > 0)
    penalty = sum(abs(weight) * scores[criterion_id] for criterion_id, weight in WEIGHTS.items() if weight < 0)
    reward = max(0.0, min(1.0, (earned - penalty) / positive_weight))
    hurdle_failures = [] if scores["R001"] == 1.0 else ["R001"]
    if hurdle_failures:
        reward = 0.0
    result = {
        "task_id": SPEC["task_id"],
        "rubric_version": "cowork-evo-rubric-v0.1",
        "judge_version": "ppt-deterministic-v1.0",
        "artifact_sha256": sorted(set(hashes)),
        "criterion_results": criteria,
        "reward": round(reward, 4),
        "passed": not hurdle_failures and reward >= 0.7,
        "hurdle_failures": hurdle_failures,
    }
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description="对PPT任务产物做确定性评分")
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    args = parser.parse_args()
    result = evaluate(args.artifact_dir)
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"task_id": result["task_id"], "reward": result["reward"], "passed": result["passed"]}, ensure_ascii=False))
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
