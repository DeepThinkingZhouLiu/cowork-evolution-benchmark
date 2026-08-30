#!/usr/bin/env python3
"""PPTArena Case 1 专用 Golden 构建器：复核并校正上游人工译文。"""

from __future__ import annotations

import argparse
import copy
import hashlib
from pathlib import Path

from pptx import Presentation


INPUT_SHA256 = "166e95dcf7883aa0ed60768e184c225851161a908f82223a7f8443f45cdd32dd"
UPSTREAM_SHA256 = "32a2adb5d145439a959d53b7dc2f24dabb54d01e37cc0cbd2c0a9b3c34dde7c2"
CORRECTIONS = {
    (4, 4): [("meal", "food"), ("a diet", "groceries")],
    (6, 2): [("classification", "conjugation")],
    (12, 5): [("To drink", "To drink water")],
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def shape_by_id(presentation: Presentation, slide_number: int, shape_id: int):
    matches = [shape for shape in presentation.slides[slide_number - 1].shapes if shape.shape_id == shape_id]
    if len(matches) != 1:
        raise ValueError(f"对象不唯一：{slide_number}/{shape_id}")
    return matches[0]


def replace_once(shape, old: str, new: str) -> None:
    count = 0
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if old in run.text:
                occurrences = run.text.count(old)
                run.text = run.text.replace(old, new)
                count += occurrences
    if count != 1:
        raise ValueError(f"替换计数不是 1：{old!r} -> {new!r}，实际 {count}")


def structure_signature(presentation: Presentation) -> tuple:
    return tuple(
        tuple((shape.shape_id, str(shape.shape_type)) for shape in slide.shapes)
        for slide in presentation.slides
    )


def build(source: Path, upstream: Path, output: Path) -> None:
    if sha256(source) != INPUT_SHA256 or sha256(upstream) != UPSTREAM_SHA256:
        raise ValueError("冻结输入或上游 Ground Truth SHA256 不匹配")
    original = Presentation(source)
    presentation = Presentation(upstream)
    if len(original.slides) != len(presentation.slides) or len(presentation.slides) != 18:
        raise ValueError("页数不再是 18")
    if structure_signature(original) != structure_signature(presentation):
        raise ValueError("上游 Ground Truth 增删了页面对象")
    # 上游在合并 run 时重写了两个法语文本框。文字和高层字体属性虽相同，
    # 但字距渲染已不再逐像素一致，因此直接恢复原始 OOXML 对象。
    for key in ((10, 2), (12, 4)):
        source_shape = shape_by_id(original, *key)
        target_shape = shape_by_id(presentation, *key)
        target_shape._element.addprevious(copy.deepcopy(source_shape._element))
        target_shape._element.getparent().remove(target_shape._element)
    for key, replacements in CORRECTIONS.items():
        shape = shape_by_id(presentation, *key)
        for old, new in replacements:
            replace_once(shape, old, new)
    coffee_cells = []
    for shape in presentation.slides[7].shapes:
        if not shape.has_table:
            continue
        for row in shape.table.rows:
            for cell in row.cells:
                if cell.text == "Cafe":
                    coffee_cells.append(cell)
    if len(coffee_cells) != 1:
        raise ValueError(f"上游 Cafe 单元格数量异常：{len(coffee_cells)}")
    replace_once(coffee_cells[0], "Cafe", "Coffee")

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    reopened = Presentation(output)
    if structure_signature(reopened) != structure_signature(original):
        raise ValueError("保存后增删了页面对象")
    expected = ["food, groceries", "verbs and their conjugation", "Coffee", "To drink water"]
    combined = "\n".join(
        cell.text if shape.has_table else shape.text
        for slide in reopened.slides
        for shape in slide.shapes
        for cell in ([cell for row in shape.table.rows for cell in row.cells] if shape.has_table else [None])
        if shape.has_table or getattr(shape, "has_text_frame", False)
    )
    if any(value not in combined for value in expected):
        raise ValueError("校正后译文缺失")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--upstream", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.source, args.upstream, args.output)
    print(f"{args.output}\nsha256={sha256(args.output)}")


if __name__ == "__main__":
    main()
