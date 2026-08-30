#!/usr/bin/env python3
"""PPTArena Case 1 专用 Judge 校准样例。"""

from __future__ import annotations

import argparse
import shutil
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches


OUTPUT_NAME = "TranslatetoEnglishButKeepallFrenchGroundTruthA.pptx"


def output_path(root: Path, case: str) -> Path:
    path = root / case / OUTPUT_NAME
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def shape_by_id(presentation: Presentation, slide_number: int, shape_id: int):
    matches = [shape for shape in presentation.slides[slide_number - 1].shapes if shape.shape_id == shape_id]
    if len(matches) != 1:
        raise ValueError(f"对象不唯一：{slide_number}/{shape_id}")
    return matches[0]


def replace_once(shape, old: str, new: str) -> None:
    count = 0
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if old in run.text:
                count += run.text.count(old)
                run.text = run.text.replace(old, new)
    if count != 1:
        raise ValueError(f"变异目标计数异常：{old!r}")


def save_mutation(golden: Path, output: Path, callback) -> None:
    presentation = Presentation(golden)
    callback(presentation)
    presentation.save(output)


def repackage(golden: Path, output: Path) -> None:
    with zipfile.ZipFile(golden) as reader, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_STORED) as writer:
        for info in reader.infolist():
            writer.writestr(info.filename, reader.read(info.filename))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--upstream", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--slide9-image", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()
    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))
    shutil.copy2(args.upstream, output_path(args.outdir, "upstream-uncorrected"))
    repackage(args.golden, output_path(args.outdir, "valid-repackage"))

    def missed_translation(prs):
        shape_by_id(prs, 8, 2).table.cell(1, 1).text = "кофе"
    save_mutation(args.golden, output_path(args.outdir, "missed-translation"), missed_translation)

    def translated_french(prs):
        shape_by_id(prs, 8, 6).table.cell(0, 0).text = "Bread"
    save_mutation(args.golden, output_path(args.outdir, "translated-french"), translated_french)

    save_mutation(args.golden, output_path(args.outdir, "wrong-conjugation"), lambda prs: replace_once(shape_by_id(prs, 6, 2), "conjugation", "classification"))
    save_mutation(args.golden, output_path(args.outdir, "wrong-water"), lambda prs: replace_once(shape_by_id(prs, 12, 5), "To drink water", "To drink"))

    def wrong_style(prs):
        for paragraph in shape_by_id(prs, 17, 2).text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(255, 0, 0)
    save_mutation(args.golden, output_path(args.outdir, "wrong-style"), wrong_style)

    def moved_target(prs):
        shape_by_id(prs, 17, 3).left += Inches(1.0)
    save_mutation(args.golden, output_path(args.outdir, "moved-target"), moved_target)

    def deleted_shape(prs):
        shape = shape_by_id(prs, 17, 3)
        shape._element.getparent().remove(shape._element)
    save_mutation(args.golden, output_path(args.outdir, "deleted-shape"), deleted_shape)

    def flattened(prs):
        slide = prs.slides[8]
        for shape in list(slide.shapes):
            shape._element.getparent().remove(shape._element)
        slide.shapes.add_picture(str(args.slide9_image), 0, 0, width=prs.slide_width, height=prs.slide_height)
    save_mutation(args.golden, output_path(args.outdir, "flattened-slide9"), flattened)

    def hardcoded(prs):
        box = prs.slides[0].shapes.add_textbox(Inches(0.1), Inches(0.1), Inches(2.0), Inches(0.2))
        box.text = "solution/Golden"
    save_mutation(args.golden, output_path(args.outdir, "hardcoded-marker"), hardcoded)

    extra = args.outdir / "extra-output"
    extra.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.golden, extra / OUTPUT_NAME)
    (extra / "debug.txt").write_text("debug\n", encoding="utf-8")


if __name__ == "__main__":
    main()
