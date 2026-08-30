#!/usr/bin/env python3
"""PPTArena Case 1 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import math
import posixpath
import re
import statistics
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from lxml import etree
from PIL import Image, ImageChops, ImageDraw, ImageStat
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "TranslatetoEnglishButKeepallFrenchGroundTruthA.pptx"
INPUT_NAME = "TranslateToEnglishButKeepallFrenchTestA_Bonjour mes amis!.pptx"
INPUT_SHA256 = "166e95dcf7883aa0ed60768e184c225851161a908f82223a7f8443f45cdd32dd"
TASK_ID = "cowork-evo/pptarena-001-kazakh-to-english-preserve-french"
RUBRIC_VERSION = "pptarena-case-001-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-001-deterministic-v0.1"
SUSPICIOUS_MARKERS = ["solution/golden", "judge-result", "artifact_dir", "pptarena-case-001", "/root/"]

EXPECTED_SHAPES = {
    (4, 4): "Le repas, une alimentation \x0b=\x0bfood, groceries",
    (5, 2): "The purpose of the lesson:\x0b\x0bTo learn new words related to food and be able to use them. To learn new verbs related to cooking.\x0b\x0bApprendre à connaître de nouveaux mots relatives une alimentation. Apprenez de nouveaux verbes relatives de la cuisine.\x0b",
    (6, 2): "Lesson plan:\x0b\x0b1. Homework check \x0b2. New words \x0b2. New verbs and their conjugation \x0b3. Practical tasks \x0b4. Feedback \x0b5. Homework\x0b",
    (10, 3): "1. in the dining hall\n\n2. at home\n\n3. in a restaurant\n\n4. in the cafe\n\n5. at school",
    (12, 5): "To eat\nTo drink water\nTo be hungry\nTo be thirsty\n",
    (12, 6): "verbs - verbs",
    (17, 2): "Feedback",
    (17, 3): "Was the lesson clear?\n\nWhat did you like and what didn't you like?\n\nMy suggestion.",
}

EXPECTED_TABLES = {
    (8, 6): ["Bread", "Salt", "Sugar ", "Butter", "Cheese", "Meat", "Sausage", "Fish", "Flour"],
    (8, 7): ["Candy", "Cookies", "Jam", "Honey", "Cake", "Sweet Bread"],
    (8, 2): ["Tea", "Coffee", "Milk", "Juice", "Lemonade"],
    (9, 2): [
        "vegetables", "cabbage", "tomato", "cucumbers", "carrots", "onion", "beetroot", "watermelon", "melon", "potatoes",
        "fruits", "apples", "pears", "grapes", "plums", "apricots", "peach", "strawberry", "raspberry", "cherry",
    ],
}

FRENCH_PREFIXES = {
    (4, 4): "Le repas, une alimentation \x0b=\x0b",
    (5, 2): "Apprendre à connaître de nouveaux mots relatives une alimentation. Apprenez de nouveaux verbes relatives de la cuisine.\x0b",
}

SPECS = {
    "R001": ("PPTX 完整且严格保留原尺寸 18 页", 7, "hurdle"),
    "R002": ("八个目标文本对象的哈萨克语英译正确", 20, "hurdle"),
    "R003": ("第 8、9 页 40 个右列单元格英译正确", 20, "hurdle"),
    "R004": ("所有法语与非目标内容完整保留", 18, "hurdle"),
    "R005": ("英译文本保留原有可见字体层级", 10, "hurdle"),
    "R006": ("对象、媒体、表格、备注和非目标几何保持", 12, "hurdle"),
    "R007": ("18 页可渲染，原生可编辑且非目标区域保持", 9, "hurdle"),
    "R008": ("无宏、损坏关系、多余输出或答案标记", 4, "positive"),
    "R009": ("无漏译错译、法语改写、样式结构破坏或压平等严重错误", -20, "penalty"),
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def safe(callback, default=""):
    try:
        return callback()
    except Exception:
        return default


def locate_source() -> Path | None:
    candidates = [Path("/app/input_files") / INPUT_NAME, Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME]
    return next((path for path in candidates if path.is_file()), None)


def shape_map(presentation: Presentation, slide_number: int) -> dict:
    return {shape.shape_id: shape for shape in presentation.slides[slide_number - 1].shapes}


def font_signature(font, include_language: bool = False) -> tuple:
    color = font.color
    values = (
        font.name,
        round(font.size.pt, 4) if font.size else None,
        font.bold,
        font.italic,
        font.underline,
        safe(lambda: str(color.type)),
        safe(lambda: str(color.rgb)),
        safe(lambda: str(color.theme_color)),
        safe(lambda: round(color.brightness, 5)),
    )
    return values + ((safe(lambda: str(font.language_id))),) if include_language else values


def character_styles(paragraph, include_language: bool = False) -> tuple:
    return tuple(style for run in paragraph.runs for style in [font_signature(run.font, include_language)] * len(run.text))


def visible_text_signature(shape) -> tuple:
    return (
        shape.text,
        tuple(
            (
                paragraph.text,
                character_styles(paragraph),
                paragraph.level,
                str(paragraph.alignment),
                str(paragraph.line_spacing),
                str(paragraph.space_before),
                str(paragraph.space_after),
            )
            for paragraph in shape.text_frame.paragraphs
        ),
        (
            shape.text_frame.margin_left,
            shape.text_frame.margin_right,
            shape.text_frame.margin_top,
            shape.text_frame.margin_bottom,
            str(shape.text_frame.vertical_anchor),
            shape.text_frame.word_wrap,
            str(shape.text_frame.auto_size),
        ),
    )


def cell_visible_signature(cell) -> tuple:
    return (
        cell.text,
        tuple((paragraph.text, character_styles(paragraph), paragraph.level, str(paragraph.alignment)) for paragraph in cell.text_frame.paragraphs),
        canonical(cell._tc.tcPr),
    )


def canonical(element) -> bytes:
    return etree.tostring(element, method="c14n") if element is not None else b""


def shape_nontext_signature(shape, ignore_geometry: bool = False) -> bytes:
    element = copy.deepcopy(shape._element)
    for body in element.xpath(".//*[local-name()='txBody']"):
        body.getparent().remove(body)
    if ignore_geometry:
        for transform in element.xpath(".//*[local-name()='xfrm']"):
            transform.getparent().remove(transform)
    return canonical(element)


def table_style_signature(shape) -> bytes:
    element = copy.deepcopy(shape._element)
    for body in element.xpath(".//*[local-name()='txBody']"):
        for paragraph in list(body.xpath("./*[local-name()='p']")):
            body.remove(paragraph)
    for transform in element.xpath(".//*[local-name()='xfrm']"):
        transform.getparent().remove(transform)
    for row in element.xpath(".//*[local-name()='tr']"):
        row.attrib.pop("h", None)
    for volatile in list(element.xpath(".//*[local-name()='modId' or local-name()='rowId']")):
        volatile.getparent().remove(volatile)
    return canonical(element)


def font_style_compatible(source_style: tuple, candidate_style: tuple) -> bool:
    source_name, candidate_name = source_style[0], candidate_style[0]
    name_ok = source_name == candidate_name or {source_name, candidate_name}.issubset({None, "Times New Roman"})
    return name_ok and source_style[1:] == candidate_style[1:]


def package_checks(path: Path) -> tuple[bool, dict]:
    evidence = {"broken_member": None, "missing_targets": [], "macros": []}
    try:
        with zipfile.ZipFile(path) as archive:
            evidence["broken_member"] = archive.testzip()
            names = set(archive.namelist())
            evidence["macros"] = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
            for rels_name in sorted(name for name in names if name.endswith(".rels")):
                root = etree.fromstring(archive.read(rels_name))
                if rels_name == "_rels/.rels":
                    owner = ""
                else:
                    prefix, filename = rels_name.split("/_rels/", 1)
                    owner = posixpath.join(prefix, filename[:-5])
                base = posixpath.dirname(owner)
                for relation in root:
                    if relation.attrib.get("TargetMode") == "External":
                        continue
                    target = relation.attrib.get("Target", "")
                    resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                    if resolved not in names:
                        evidence["missing_targets"].append(f"{rels_name} -> {resolved}")
        passed = evidence["broken_member"] is None and not evidence["missing_targets"] and not evidence["macros"]
        return passed, evidence
    except Exception as exc:
        evidence["exception"] = str(exc)
        return False, evidence


def media_hashes(path: Path) -> Counter[str]:
    with zipfile.ZipFile(path) as archive:
        return Counter(hashlib.sha256(archive.read(name)).hexdigest() for name in archive.namelist() if name.startswith("ppt/media/"))


def notes_signature(slide) -> tuple[str, ...]:
    try:
        return tuple(shape.text for shape in slide.notes_slide.shapes if getattr(shape, "has_text_frame", False))
    except Exception:
        return ()


def geometry(shape) -> tuple[int, int, int, int]:
    return shape.left, shape.top, shape.width, shape.height


def geometry_allowed(source_shape, candidate_shape, target: bool, presentation: Presentation) -> bool:
    if not target:
        return geometry(source_shape) == geometry(candidate_shape)
    source_values = geometry(source_shape)
    candidate_values = geometry(candidate_shape)
    tolerances = (0.65 * 914400, 0.65 * 914400, 0.65 * 914400, 1.5 * 914400)
    close = all(abs(a - b) <= tolerance for a, b, tolerance in zip(source_values, candidate_values, tolerances, strict=True))
    in_bounds = candidate_shape.left >= 0 and candidate_shape.top >= 0 and candidate_shape.left + candidate_shape.width <= presentation.slide_width and candidate_shape.top + candidate_shape.height <= presentation.slide_height
    return close and in_bounds


def translated_style_ok(source_shape, candidate_shape, key: tuple[int, int]) -> bool:
    candidate_runs = [run for paragraph in candidate_shape.text_frame.paragraphs for run in paragraph.runs if run.text]
    if not candidate_runs:
        return False
    allowed = {font_signature(run.font) for paragraph in source_shape.text_frame.paragraphs for run in paragraph.runs if run.text}
    if not all(font_signature(run.font) in allowed for run in candidate_runs):
        return False
    text = candidate_shape.text
    char_styles = tuple(style for paragraph in candidate_shape.text_frame.paragraphs for style in character_styles(paragraph))
    if len(char_styles) != len(text.replace("\n", "")) and "\n" in text:
        pass
    if key in {(4, 4), (17, 2)}:
        return all(run.font.bold is True and run.font.italic is True for run in candidate_runs)
    if key in {(5, 2), (6, 2)}:
        first_break = text.find("\x0b")
        if first_break < 0:
            return False
        cursor = 0
        header_ok = body_ok = True
        for paragraph in candidate_shape.text_frame.paragraphs:
            for run in paragraph.runs:
                start, end = cursor, cursor + len(run.text)
                if end <= first_break and run.text:
                    header_ok = header_ok and run.font.bold is True and run.font.italic is True
                elif run.text:
                    body_ok = body_ok and run.font.bold in (None, False) and run.font.italic is True
                cursor = end
        return header_ok and body_ok
    if key == (12, 6):
        return all(run.font.size and abs(run.font.size.pt - 28) < 0.01 and str(run.font.color.rgb) == "FF0000" for run in candidate_runs)
    if key == (17, 3):
        return all(run.font.size and abs(run.font.size.pt - 24) < 0.01 and run.font.bold in (None, False) for run in candidate_runs)
    return True


def render(path: Path, outdir: Path) -> tuple[list[Path], int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "profile"
    profile.mkdir()
    completed = subprocess.run(
        ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(path)],
        stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300, check=False,
    )
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return [], 0, completed.stdout[-500:]
    info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=30, check=False).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    pages = int(match.group(1)) if match else 0
    prefix = outdir / "slide"
    raster = subprocess.run(["pdftoppm", "-png", "-r", "72", str(pdf), str(prefix)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300, check=False)
    images = sorted(outdir.glob("slide-*.png")) if raster.returncode == 0 else []
    return images, pages, f"soffice={completed.returncode}, raster={raster.returncode}"


def image_box(shape, image: Image.Image, presentation: Presentation) -> tuple[int, int, int, int]:
    pixels_per_inch = image.width / (presentation.slide_width / 914400)
    padding = int(math.ceil(1.1 * pixels_per_inch))
    return (
        max(0, int(shape.left * image.width / presentation.slide_width) - padding),
        max(0, int(shape.top * image.height / presentation.slide_height) - padding),
        min(image.width, int(math.ceil((shape.left + shape.width) * image.width / presentation.slide_width)) + padding),
        min(image.height, int(math.ceil((shape.top + shape.height) * image.height / presentation.slide_height)) + padding),
    )


def render_similarity(source_images: list[Path], candidate_images: list[Path], source: Presentation, candidate: Presentation) -> float:
    targets_by_slide: dict[int, set[int]] = {}
    for slide_number, shape_id in EXPECTED_SHAPES:
        targets_by_slide.setdefault(slide_number, set()).add(shape_id)
    for slide_number, shape_id in EXPECTED_TABLES:
        targets_by_slide.setdefault(slide_number, set()).add(shape_id)
    similarities = []
    for slide_number, (source_path, candidate_path) in enumerate(zip(source_images, candidate_images, strict=True), 1):
        with Image.open(source_path) as source_raw, Image.open(candidate_path) as candidate_raw:
            first, second = source_raw.convert("RGB"), candidate_raw.convert("RGB")
            if first.size != second.size:
                similarities.append(0.0)
                continue
            mask = Image.new("L", first.size, 255)
            draw = ImageDraw.Draw(mask)
            for shape_id in targets_by_slide.get(slide_number, set()):
                source_shape = shape_map(source, slide_number).get(shape_id)
                candidate_shape = shape_map(candidate, slide_number).get(shape_id)
                if source_shape is not None:
                    draw.rectangle(image_box(source_shape, first, source), fill=0)
                if candidate_shape is not None:
                    draw.rectangle(image_box(candidate_shape, second, candidate), fill=0)
            difference = ImageChops.difference(first, second)
            difference.paste((0, 0, 0), mask=ImageChops.invert(mask))
            similarities.append(max(0.0, 1.0 - statistics.mean(ImageStat.Stat(difference).rms) / 255.0))
    return min(similarities, default=0.0)


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {"criterion_id": identifier, "score": 1.0 if passed else 0.0, "evidence": evidence, "method": "deterministic", "raw": {"description": description, "weight": weight, "type": kind, "passed": bool(passed)}}


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [item["criterion_id"] for item in criteria if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0 and item["raw"]["passed"])
    penalty = sum(abs(item["raw"]["weight"]) for item in criteria if item["raw"]["weight"] < 0 and not item["raw"]["passed"])
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [sha256(path)] if path.is_file() else [], "criterion_results": criteria, "reward": reward, "passed": reward >= 0.8 and not hurdles, "hurdle_failures": hurdles}


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        return make_result(path, [result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配") for identifier in SPECS])
    try:
        source, candidate = Presentation(source_path), Presentation(path)
    except Exception as exc:
        return make_result(path, [result_item(identifier, False, f"PPTX 打开失败：{exc}") for identifier in SPECS])

    criteria = []
    package_passed, package_evidence = package_checks(path)
    size_passed = len(source.slides) == len(candidate.slides) == 18 and (source.slide_width, source.slide_height) == (candidate.slide_width, candidate.slide_height)
    criteria.append(result_item("R001", package_passed and size_passed, f"页数={len(candidate.slides)}，尺寸一致={size_passed}，包={json.dumps(package_evidence, ensure_ascii=False)}"))

    shape_translation = []
    shape_translation_passed = True
    style_passed = True
    for key, expected in EXPECTED_SHAPES.items():
        source_shape = shape_map(source, key[0]).get(key[1])
        candidate_shape = shape_map(candidate, key[0]).get(key[1])
        text_ok = candidate_shape is not None and getattr(candidate_shape, "has_text_frame", False) and candidate_shape.text == expected
        style_ok = bool(text_ok and translated_style_ok(source_shape, candidate_shape, key))
        shape_translation_passed = shape_translation_passed and text_ok
        style_passed = style_passed and style_ok
        shape_translation.append({"shape": key, "text": text_ok, "style": style_ok})
    criteria.append(result_item("R002", shape_translation_passed, json.dumps(shape_translation, ensure_ascii=False)))

    table_translation_passed = True
    table_style_passed = True
    table_details = []
    for key, expected_values in EXPECTED_TABLES.items():
        source_shape = shape_map(source, key[0]).get(key[1])
        candidate_shape = shape_map(candidate, key[0]).get(key[1])
        actual_values = []
        styles_ok = candidate_shape is not None and candidate_shape.has_table and len(candidate_shape.table.rows) == len(expected_values) and len(candidate_shape.table.columns) == 2
        if styles_ok:
            for row_index, expected in enumerate(expected_values):
                source_cell = source_shape.table.cell(row_index, 1)
                candidate_cell = candidate_shape.table.cell(row_index, 1)
                actual_values.append(candidate_cell.text)
                source_styles = {font_signature(run.font) for paragraph in source_cell.text_frame.paragraphs for run in paragraph.runs if run.text}
                candidate_styles = {font_signature(run.font) for paragraph in candidate_cell.text_frame.paragraphs for run in paragraph.runs if run.text}
                styles_ok = styles_ok and bool(candidate_styles) and all(any(font_style_compatible(source_style, candidate_style) for source_style in source_styles) for candidate_style in candidate_styles)
        values_ok = actual_values == expected_values
        table_translation_passed = table_translation_passed and values_ok
        table_style_passed = table_style_passed and styles_ok
        table_details.append({"table": key, "rows": len(actual_values), "text": values_ok, "style": styles_ok})
    style_passed = style_passed and table_style_passed
    criteria.append(result_item("R003", table_translation_passed, json.dumps(table_details, ensure_ascii=False)))

    preservation_passed = True
    preservation_details = []
    target_shape_keys = set(EXPECTED_SHAPES)
    target_table_keys = set(EXPECTED_TABLES)
    structure_passed = len(source.slides) == len(candidate.slides)
    layout_passed = True
    table_styles_passed = True
    for slide_number, (source_slide, candidate_slide) in enumerate(zip(source.slides, candidate.slides, strict=False), 1):
        source_shapes, candidate_shapes = shape_map(source, slide_number), shape_map(candidate, slide_number)
        structure_passed = structure_passed and set(source_shapes) == set(candidate_shapes)
        if set(source_shapes) != set(candidate_shapes):
            preservation_passed = layout_passed = False
            continue
        for shape_id, source_shape in source_shapes.items():
            candidate_shape = candidate_shapes[shape_id]
            key = (slide_number, shape_id)
            is_target = key in target_shape_keys or key in target_table_keys
            structure_passed = structure_passed and source_shape.shape_type == candidate_shape.shape_type
            layout_passed = layout_passed and geometry_allowed(source_shape, candidate_shape, is_target, candidate)
            if source_shape.has_table:
                if not candidate_shape.has_table:
                    preservation_passed = table_styles_passed = False
                    continue
                table_styles_passed = table_styles_passed and table_style_signature(source_shape) == table_style_signature(candidate_shape)
                if key in target_table_keys:
                    for row_index in range(len(source_shape.table.rows)):
                        preservation_passed = preservation_passed and cell_visible_signature(source_shape.table.cell(row_index, 0)) == cell_visible_signature(candidate_shape.table.cell(row_index, 0))
                else:
                    for row_index in range(len(source_shape.table.rows)):
                        for column_index in range(len(source_shape.table.columns)):
                            preservation_passed = preservation_passed and cell_visible_signature(source_shape.table.cell(row_index, column_index)) == cell_visible_signature(candidate_shape.table.cell(row_index, column_index))
            elif getattr(source_shape, "has_text_frame", False):
                if key not in target_shape_keys:
                    preservation_passed = preservation_passed and visible_text_signature(source_shape) == visible_text_signature(candidate_shape)
                else:
                    preservation_passed = preservation_passed and shape_nontext_signature(source_shape, ignore_geometry=True) == shape_nontext_signature(candidate_shape, ignore_geometry=True)
            else:
                preservation_passed = preservation_passed and shape_nontext_signature(source_shape, ignore_geometry=is_target) == shape_nontext_signature(candidate_shape, ignore_geometry=is_target)
        preservation_passed = preservation_passed and notes_signature(source_slide) == notes_signature(candidate_slide)
    try:
        preservation_passed = preservation_passed and media_hashes(source_path) == media_hashes(path)
    except Exception:
        preservation_passed = False
    preservation_details.append({"objects": structure_passed, "layout": layout_passed, "table_styles": table_styles_passed, "media": safe(lambda: media_hashes(source_path) == media_hashes(path), False)})

    french_passed = preservation_passed
    for key, invariant in FRENCH_PREFIXES.items():
        candidate_text = shape_map(candidate, key[0])[key[1]].text
        french_passed = french_passed and invariant in candidate_text
    criteria.append(result_item("R004", french_passed, json.dumps(preservation_details, ensure_ascii=False)))
    criteria.append(result_item("R005", style_passed, f"目标对象样式={shape_translation}，表格右列样式={table_style_passed}"))

    source_full_slide_pictures = {
        (slide_number, shape.shape_id)
        for slide_number, slide in enumerate(source.slides, 1)
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.width * shape.height >= source.slide_width * source.slide_height * 0.8
    }
    candidate_full_slide_pictures = {
        (slide_number, shape.shape_id)
        for slide_number, slide in enumerate(candidate.slides, 1)
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.width * shape.height >= candidate.slide_width * candidate.slide_height * 0.8
    }
    extra_full_slide_pictures = sorted(candidate_full_slide_pictures - source_full_slide_pictures)
    layout_and_structure = structure_passed and layout_passed and table_styles_passed and not extra_full_slide_pictures
    criteria.append(result_item("R006", layout_and_structure, f"对象={structure_passed}，几何={layout_passed}，表格样式={table_styles_passed}，新增整页图={extra_full_slide_pictures}"))

    render_passed = False
    render_evidence = "未渲染"
    pre_render_passed = shape_translation_passed and table_translation_passed and french_passed and style_passed and layout_and_structure
    if not pre_render_passed:
        render_evidence = "前置内容、保持、样式或结构门槛已失败，不再浪费容器时间渲染错误交付"
    else:
        try:
            with tempfile.TemporaryDirectory(prefix="pptarena001-render-") as temporary:
                root = Path(temporary)
                source_images, source_pages, source_log = render(source_path, root / "source")
                candidate_images, candidate_pages, candidate_log = render(path, root / "candidate")
                if source_pages == candidate_pages == 18 and len(source_images) == len(candidate_images) == 18:
                    similarity = render_similarity(source_images, candidate_images, source, candidate)
                    render_passed = similarity >= 0.995 and structure_passed and not extra_full_slide_pictures
                    render_evidence = f"非目标区域最低相似度={similarity:.6f}"
                else:
                    render_evidence = f"页数={source_pages}/{candidate_pages}，{source_log}，{candidate_log}"
        except Exception as exc:
            render_evidence = f"渲染异常：{exc}"
    criteria.append(result_item("R007", render_passed, render_evidence))

    delivered = sorted(str(item.relative_to(artifact_dir)) for item in artifact_dir.rglob("*") if item.is_file())
    combined_text = "\n".join(shape.text for slide in candidate.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False)).lower()
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker in combined_text]
    hygiene_passed = delivered == [OUTPUT_NAME] and package_passed and not markers
    criteria.append(result_item("R008", hygiene_passed, f"文件={delivered}，标记={markers}"))

    severe = not (shape_translation_passed and table_translation_passed and french_passed and style_passed and layout_and_structure and render_passed)
    criteria.append(result_item("R009", not severe, f"严重错误={severe}"))
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    path = args.artifact_dir / OUTPUT_NAME
    if path.is_file():
        result = evaluate(path, args.artifact_dir)
    else:
        result = make_result(path, [result_item(identifier, False, f"缺失指定交付物：{OUTPUT_NAME}") for identifier in SPECS])
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{result['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
