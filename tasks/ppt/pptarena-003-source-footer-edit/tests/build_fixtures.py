#!/usr/bin/env python3
"""只为 PPTArena Case 3 构造可重放的 Judge 负例和中等例。"""

from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt


OUTPUT_NAME = "VerifyMe_Q1_2025_with_source.pptx"
CITATION = "Source: www.VerifyMe.com"


def output_path(root: Path, case: str) -> Path:
    target = root / case / OUTPUT_NAME
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def add_citation(source: Path, output: Path, slide_index: int, text: str, top: float) -> None:
    presentation = Presentation(source)
    box = presentation.slides[slide_index].shapes.add_textbox(
        Inches(0.5), Inches(top), Inches(12.333), Inches(0.25)
    )
    box.text = text
    for paragraph in box.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(10)
    presentation.save(output)


def changed_original_text(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    replaced = False
    for shape in presentation.slides[4].shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if "22.6%" in run.text:
                    run.text = run.text.replace("22.6%", "99.9%")
                    replaced = True
                    break
            if replaced:
                break
        if replaced:
            break
    if not replaced:
        raise ValueError("未找到待变异的 22.6%")
    presentation.save(output)


def flattened_slide(source: Path, slide_image: Path, output: Path) -> None:
    presentation = Presentation(source)
    slide = presentation.slides[4]
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    slide.shapes.add_picture(
        str(slide_image), 0, 0, width=presentation.slide_width, height=presentation.slide_height
    )
    presentation.save(output)


def duplicate_citation(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    box = presentation.slides[4].shapes.add_textbox(
        Inches(6.6), Inches(6.7), Inches(5.8), Inches(0.25)
    )
    box.text = CITATION
    for paragraph in box.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(10)
    presentation.save(output)


def hardcoded_marker(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    box = presentation.slides[4].shapes.add_textbox(
        Inches(0.1), Inches(7.6), Inches(2.0), Inches(0.2)
    )
    box.text = "solution/Golden"
    presentation.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--slide5-image", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()

    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))
    add_citation(args.source, output_path(args.outdir, "wrong-text"), 4, "Source: www.VerifyMe.co", 7.0)
    add_citation(args.source, output_path(args.outdir, "wrong-slide"), 3, CITATION, 7.0)
    add_citation(args.source, output_path(args.outdir, "off-slide"), 4, CITATION, 7.55)
    changed_original_text(args.golden, output_path(args.outdir, "changed-original-text"))
    flattened_slide(args.source, args.slide5_image, output_path(args.outdir, "flattened"))
    duplicate_citation(args.golden, output_path(args.outdir, "duplicate"))
    hardcoded_marker(args.golden, output_path(args.outdir, "hardcoded-marker"))


if __name__ == "__main__":
    main()
