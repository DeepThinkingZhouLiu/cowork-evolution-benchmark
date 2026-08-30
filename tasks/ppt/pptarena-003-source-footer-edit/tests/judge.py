#!/usr/bin/env python3
"""PPTArena Case 3 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from PIL import Image, ImageChops
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "VerifyMe_Q1_2025_with_source.pptx"
INPUT_NAME = "VerifyMe_Q1_2025.pptx"
TASK_ID = "cowork-evo/pptarena-003-source-footer-edit"
RUBRIC_VERSION = "pptarena-case-003-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-003-deterministic-v0.1"
CITATION = "Source: www.VerifyMe.com"
CITATION_NORMALIZED = "Source:www.VerifyMe.com"
EMU_PER_INCH = 914_400

SPECS = {
    "R001": ("指定 PPTX 可打开，严格保留 13 页与页面尺寸", 5, "hurdle"),
    "R002": ("第 5 页恰有一个文字精确的原生可编辑引用", 8, "hurdle"),
    "R003": ("引用位于第 5 页分隔线下方的页脚区域且完整在页内", 4, "positive"),
    "R004": ("引用字号与文本框高度可读，渲染后可提取", 4, "positive"),
    "R005": ("全部原有文字按原页完整保留", 8, "hurdle"),
    "R006": ("除第 5 页外的 12 页渲染与冻结输入高度一致", 6, "hurdle"),
    "R007": ("第 5 页分隔线上方的主内容渲染与冻结输入高度一致", 6, "hurdle"),
    "R008": ("原有图片媒体保持，没有新增整页压平图片", 5, "positive"),
    "R009": ("指定引用只出现在第 5 页一次", 3, "positive"),
    "R010": ("LibreOffice 可渲染 13 页，PDF 第 5 页可提取完整引用", 4, "positive"),
    "R011": ("页数、其他页、原文、引用次数或可编辑性发生严重错误", -15, "penalty"),
}


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", "", value or "")


def locate_source() -> Path | None:
    candidates = [
        Path("/app/input_files") / INPUT_NAME,
        Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME,
    ]
    return next((path for path in candidates if path.is_file()), None)


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def slide_text_counter(slide) -> Counter[str]:
    values: Counter[str] = Counter()
    for shape in iter_shapes(slide.shapes):
        if getattr(shape, "has_text_frame", False):
            text = normalize_text(shape.text)
            if text:
                values[text] += 1
    return values


def slide_picture_counter(slide) -> Counter[str]:
    values: Counter[str] = Counter()
    for shape in iter_shapes(slide.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            values[hashlib.sha256(shape.image.blob).hexdigest()] += 1
    return values


def citation_font_sizes(shape) -> list[float]:
    sizes: list[float] = []
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size is not None:
                sizes.append(float(run.font.size.pt))
    for raw in shape._element.xpath(
        './/*[local-name()="rPr" or local-name()="defRPr" or local-name()="endParaRPr"]/@sz'
    ):
        try:
            sizes.append(int(raw) / 100.0)
        except (TypeError, ValueError):
            continue
    return sizes


def package_ok(path: Path) -> tuple[bool, str]:
    try:
        with zipfile.ZipFile(path) as archive:
            broken = archive.testzip()
            required = {"[Content_Types].xml", "ppt/presentation.xml", "ppt/slides/slide5.xml"}
            missing = sorted(required - set(archive.namelist()))
        return broken is None and not missing, f"损坏成员={broken}，缺失={missing}"
    except Exception as exc:
        return False, f"ZIP 检查失败：{exc}"


def render_presentation(path: Path, outdir: Path) -> tuple[list[Path], str, int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "lo-profile"
    profile.mkdir()
    completed = subprocess.run(
        [
            "soffice",
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(outdir),
            str(path),
        ],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=180,
        check=False,
    )
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return [], "", 0, f"转 PDF 失败：{completed.stdout[-500:]}"

    info = subprocess.run(
        ["pdfinfo", str(pdf)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=30,
        check=False,
    ).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    pages = int(match.group(1)) if match else 0

    prefix = outdir / "slide"
    raster = subprocess.run(
        ["pdftoppm", "-png", "-r", "120", str(pdf), str(prefix)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    images = sorted(outdir.glob("slide-*.png")) if raster.returncode == 0 else []

    page5_text = outdir / "page5.txt"
    extracted = subprocess.run(
        ["pdftotext", "-f", "5", "-l", "5", str(pdf), str(page5_text)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=30,
        check=False,
    )
    text = page5_text.read_text(encoding="utf-8", errors="replace") if extracted.returncode == 0 and page5_text.is_file() else ""
    evidence = f"页数={pages}，渲染图={len(images)}，pdftotext={extracted.returncode}"
    return images, text, pages, evidence


def image_similarity(first: Path, second: Path, crop_bottom: int | None = None) -> float:
    with Image.open(first) as left_image, Image.open(second) as right_image:
        left = left_image.convert("RGB")
        right = right_image.convert("RGB")
        if left.size != right.size:
            return 0.0
        if crop_bottom is not None:
            bottom = max(1, min(crop_bottom, left.height))
            left = left.crop((0, 0, left.width, bottom))
            right = right.crop((0, 0, right.width, bottom))
        difference = ImageChops.difference(left, right)
        histogram = difference.histogram()
        squared = sum(count * ((index % 256) ** 2) for index, count in enumerate(histogram))
        rms = math.sqrt(squared / (left.width * left.height * 3))
        return max(0.0, 1.0 - rms / 255.0)


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {
        "criterion_id": identifier,
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
        "raw": {
            "description": description,
            "weight": weight,
            "type": kind,
            "passed": bool(passed),
        },
    }


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [
        item["criterion_id"]
        for item in criteria
        if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]
    ]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(
        item["raw"]["weight"]
        for item in criteria
        if item["raw"]["weight"] > 0 and item["raw"]["passed"]
    )
    penalty = sum(
        abs(item["raw"]["weight"])
        for item in criteria
        if item["raw"]["weight"] < 0 and not item["raw"]["passed"]
    )
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [hashlib.sha256(path.read_bytes()).hexdigest()] if path.is_file() else [],
        "criterion_results": criteria,
        "reward": reward,
        "passed": reward >= 0.8 and not hurdles,
        "hurdle_failures": hurdles,
    }


def evaluate(path: Path) -> dict:
    source_path = locate_source()
    if source_path is None:
        criteria = [result_item(identifier, False, "冻结输入不可用") for identifier in SPECS]
        return make_result(path, criteria)

    try:
        candidate = Presentation(path)
        source = Presentation(source_path)
    except Exception as exc:
        criteria = [result_item(identifier, False, f"PPTX 打开失败：{exc}") for identifier in SPECS]
        return make_result(path, criteria)

    criteria: list[dict] = []
    package_passed, package_evidence = package_ok(path)
    size_passed = (
        len(candidate.slides) == 13
        and abs(candidate.slide_width / EMU_PER_INCH - 13.333333) <= 0.002
        and abs(candidate.slide_height / EMU_PER_INCH - 7.5) <= 0.002
    )
    criteria.append(
        result_item(
            "R001",
            package_passed and size_passed,
            f"{package_evidence}，页数={len(candidate.slides)}，页面={candidate.slide_width/EMU_PER_INCH:.4f}×{candidate.slide_height/EMU_PER_INCH:.4f}英寸",
        )
    )

    citation_shapes = []
    if len(candidate.slides) >= 5:
        citation_shapes = [
            shape
            for shape in iter_shapes(candidate.slides[4].shapes)
            if getattr(shape, "has_text_frame", False) and normalize_text(shape.text) == CITATION_NORMALIZED
        ]
    citation_passed = len(citation_shapes) == 1
    criteria.append(result_item("R002", citation_passed, f"第5页精确原生文本形状数={len(citation_shapes)}"))

    footer_passed = False
    font_passed = False
    footer_evidence = "未找到唯一引用形状"
    font_evidence = "未找到唯一引用形状"
    if citation_passed:
        shape = citation_shapes[0]
        left = shape.left / EMU_PER_INCH
        top = shape.top / EMU_PER_INCH
        width = shape.width / EMU_PER_INCH
        height = shape.height / EMU_PER_INCH
        footer_passed = top >= 6.38 and top + height <= 7.505 and left >= -0.01 and left + width <= 13.343
        footer_evidence = f"形状=(x={left:.3f}, y={top:.3f}, w={width:.3f}, h={height:.3f})英寸"
        font_sizes = citation_font_sizes(shape)
        explicit_ok = bool(font_sizes) and min(font_sizes) >= 8.0 and max(font_sizes) <= 16.0
        implicit_ok = not font_sizes and 0.16 <= height <= 0.65
        font_passed = explicit_ok or implicit_ok
        font_evidence = f"显式字号={font_sizes}，文本框高={height:.3f}英寸"
    criteria.append(result_item("R003", footer_passed, footer_evidence))

    source_texts = [slide_text_counter(slide) for slide in source.slides]
    candidate_texts = [slide_text_counter(slide) for slide in candidate.slides]
    text_mismatches: list[int] = []
    if len(candidate_texts) == len(source_texts):
        adjusted = [Counter(counter) for counter in candidate_texts]
        if adjusted[4][CITATION_NORMALIZED] > 0:
            adjusted[4][CITATION_NORMALIZED] -= 1
            if adjusted[4][CITATION_NORMALIZED] == 0:
                del adjusted[4][CITATION_NORMALIZED]
        text_mismatches = [index + 1 for index, (a, b) in enumerate(zip(adjusted, source_texts)) if a != b]
    else:
        text_mismatches = list(range(1, max(len(candidate_texts), len(source_texts)) + 1))
    text_passed = not text_mismatches
    criteria.append(result_item("R005", text_passed, f"原有文本不一致页={text_mismatches}"))

    source_pictures = [slide_picture_counter(slide) for slide in source.slides]
    candidate_pictures = [slide_picture_counter(slide) for slide in candidate.slides]
    picture_mismatches = []
    if len(source_pictures) == len(candidate_pictures):
        picture_mismatches = [
            index + 1
            for index, (actual, expected) in enumerate(zip(candidate_pictures, source_pictures))
            if actual != expected
        ]
    else:
        picture_mismatches = list(range(1, max(len(source_pictures), len(candidate_pictures)) + 1))
    full_slide_new = []
    for slide_index, slide in enumerate(candidate.slides, start=1):
        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            coverage = (shape.width * shape.height) / (candidate.slide_width * candidate.slide_height)
            digest = hashlib.sha256(shape.image.blob).hexdigest()
            source_counter = source_pictures[slide_index - 1] if slide_index <= len(source_pictures) else Counter()
            if coverage >= 0.9 and source_counter[digest] == 0:
                full_slide_new.append(slide_index)
    picture_passed = not picture_mismatches and not full_slide_new
    criteria.append(
        result_item(
            "R008",
            picture_passed,
            f"图片集不一致页={picture_mismatches}，新增整页图片页={full_slide_new}",
        )
    )

    citation_occurrences = []
    for slide_index, slide in enumerate(candidate.slides, start=1):
        for shape in iter_shapes(slide.shapes):
            if getattr(shape, "has_text_frame", False) and CITATION_NORMALIZED in normalize_text(shape.text):
                citation_occurrences.append(slide_index)
    occurrence_passed = citation_occurrences == [5]
    criteria.append(result_item("R009", occurrence_passed, f"引用出现页={citation_occurrences}"))

    source_images: list[Path] = []
    candidate_images: list[Path] = []
    source_render_text = ""
    candidate_render_text = ""
    source_render_pages = 0
    candidate_render_pages = 0
    render_evidence = ""
    with tempfile.TemporaryDirectory(prefix="pptarena003-judge-") as temp_dir:
        root = Path(temp_dir)
        source_images, source_render_text, source_render_pages, source_note = render_presentation(source_path, root / "source")
        candidate_images, candidate_render_text, candidate_render_pages, candidate_note = render_presentation(path, root / "candidate")
        render_evidence = f"源[{source_note}]；候选[{candidate_note}]"

        non_target_scores: list[float] = []
        if len(source_images) == 13 and len(candidate_images) == 13:
            for index in range(13):
                if index != 4:
                    non_target_scores.append(image_similarity(source_images[index], candidate_images[index]))
        non_target_min = min(non_target_scores) if non_target_scores else 0.0
        non_target_passed = len(non_target_scores) == 12 and non_target_min >= 0.999
        criteria.append(result_item("R006", non_target_passed, f"12 页最低像素相似度={non_target_min:.6f}"))

        slide5_score = 0.0
        if len(source_images) == 13 and len(candidate_images) == 13:
            with Image.open(source_images[4]) as image:
                crop_bottom = round(image.height * 6.38 / 7.5)
            slide5_score = image_similarity(source_images[4], candidate_images[4], crop_bottom=crop_bottom)
        slide5_passed = slide5_score >= 0.999
        criteria.append(result_item("R007", slide5_passed, f"第5页主内容像素相似度={slide5_score:.6f}"))

    rendered_citation = CITATION_NORMALIZED in normalize_text(candidate_render_text)
    render_passed = candidate_render_pages == 13 and len(candidate_images) == 13 and rendered_citation
    criteria.insert(3, result_item("R004", font_passed and rendered_citation, f"{font_evidence}，PDF可提取={rendered_citation}"))
    criteria.append(result_item("R010", render_passed, render_evidence + f"，引用可提取={rendered_citation}"))

    non_target_item = next(item for item in criteria if item["criterion_id"] == "R006")
    slide5_item = next(item for item in criteria if item["criterion_id"] == "R007")
    severe = (
        len(candidate.slides) != 13
        or not text_passed
        or not non_target_item["raw"]["passed"]
        or not slide5_item["raw"]["passed"]
        or len(citation_occurrences) != 1
        or bool(full_slide_new)
    )
    criteria.append(result_item("R011", not severe, f"严重错误={severe}"))

    order = {identifier: index for index, identifier in enumerate(SPECS)}
    criteria.sort(key=lambda item: order[item["criterion_id"]])
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()

    output = args.artifact_dir / OUTPUT_NAME
    if output.is_file():
        evaluation = evaluate(output)
    else:
        evaluation = {
            "task_id": TASK_ID,
            "rubric_version": RUBRIC_VERSION,
            "judge_version": JUDGE_VERSION,
            "artifact_sha256": [],
            "criterion_results": [],
            "reward": 0.0,
            "passed": False,
            "hurdle_failures": ["missing_output"],
        }

    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
