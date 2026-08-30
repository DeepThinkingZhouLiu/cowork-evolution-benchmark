#!/usr/bin/env python3
"""只为 PPTArena Case 4 构造可重放的 Judge 负例和中等例。"""

from __future__ import annotations

import argparse
import re
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt


OUTPUT_NAME = "FontSizePlus2Slide8EditedA.pptx"
TARGET_RULES = {
    "Foranobjectoutsidethefocalpoint,concavemirrorscreateimageswhichare": [22.0],
    "Inverted(upsidedown)": [26.0, 20.0],
    "Real(formedbyintersectionofactualrays)": [26.0, 20.0],
    "Objectdistance(s)": [26.0],
    "Focallengthiscalledpositive": [22.0],
    "Canbereducedorenlarged": [26.0],
    "Imagedistance(s')": [26.0],
    "Objectheight(h)": [26.0],
    "Imageheight(h')": [26.0],
}
TITLE = "ConcaveMirrors:Summary(forobjectlocatedoutsideF)"
LABELS = {"C", "F", "h", "h'", "s", "s'"}


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", "", value or "")


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def output_path(root: Path, case: str) -> Path:
    target = root / case / OUTPUT_NAME
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def find_unique(slide, key: str):
    matches = [
        shape
        for shape in iter_shapes(slide.shapes)
        if getattr(shape, "has_text_frame", False) and normalize_text(shape.text) == key
    ]
    if len(matches) != 1:
        raise ValueError(f"文本对象应唯一：{key}，实际={len(matches)}")
    return matches[0]


def set_paragraph_sizes(shape, sizes: list[float]) -> None:
    paragraphs = [
        paragraph
        for paragraph in shape.text_frame.paragraphs
        if any(run.text for run in paragraph.runs)
    ]
    if len(paragraphs) != len(sizes):
        raise ValueError(f"段落数量不匹配：{shape.text!r}")
    for paragraph, size in zip(paragraphs, sizes):
        for run in paragraph.runs:
            if run.text:
                run.font.size = Pt(size)


def wrong_amount(source: Path, output: Path) -> None:
    presentation = Presentation(source)
    slide = presentation.slides[7]
    for key, final_sizes in TARGET_RULES.items():
        set_paragraph_sizes(find_unique(slide, key), [size - 1.0 for size in final_sizes])
    presentation.save(output)


def incomplete(source: Path, output: Path) -> None:
    presentation = Presentation(source)
    slide = presentation.slides[7]
    key = "Objectheight(h)"
    set_paragraph_sizes(find_unique(slide, key), TARGET_RULES[key])
    presentation.save(output)


def changed_title(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    title = find_unique(presentation.slides[7], TITLE)
    for paragraph in title.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.text:
                run.font.size = Pt(52)
    presentation.save(output)


def changed_labels(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    changed = 0
    for shape in iter_shapes(presentation.slides[7].shapes):
        if not getattr(shape, "has_text_frame", False) or normalize_text(shape.text) not in LABELS:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.text:
                    run.font.size = Pt(20)
                    changed += 1
    if changed != 6:
        raise ValueError(f"应修改 6 个光线图标签，实际={changed}")
    presentation.save(output)


def changed_other_slide(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    for shape in iter_shapes(presentation.slides[0].shapes):
        if not getattr(shape, "has_text_frame", False):
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.text:
                    run.text += " X"
                    presentation.save(output)
                    return
    raise ValueError("第 1 页未找到可变异文字")


def wrong_style(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    shape = find_unique(presentation.slides[7], "Objectheight(h)")
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.text:
                run.font.color.rgb = RGBColor(255, 0, 0)
                presentation.save(output)
                return
    raise ValueError("未找到可变异文字运行")


def moved_target(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    shape = find_unique(
        presentation.slides[7],
        "Foranobjectoutsidethefocalpoint,concavemirrorscreateimageswhichare",
    )
    shape.top = Inches(7.2)
    presentation.save(output)


def duplicate_target(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    slide = presentation.slides[7]
    box = slide.shapes.add_textbox(Inches(0.1), Inches(0.1), Inches(5.0), Inches(0.5))
    box.text = "Object height (h)"
    for paragraph in box.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(26)
    presentation.save(output)


def hardcoded_marker(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    box = presentation.slides[7].shapes.add_textbox(
        Inches(0.1), Inches(7.6), Inches(2.0), Inches(0.2)
    )
    box.text = "solution/Golden"
    presentation.save(output)


def flattened_slide(source: Path, slide_image: Path, output: Path) -> None:
    presentation = Presentation(source)
    slide = presentation.slides[7]
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    slide.shapes.add_picture(
        str(slide_image), 0, 0, width=presentation.slide_width, height=presentation.slide_height
    )
    presentation.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--slide8-image", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()

    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))
    wrong_amount(args.source, output_path(args.outdir, "wrong-amount"))
    incomplete(args.source, output_path(args.outdir, "incomplete"))
    changed_title(args.golden, output_path(args.outdir, "changed-title"))
    changed_labels(args.golden, output_path(args.outdir, "changed-labels"))
    changed_other_slide(args.golden, output_path(args.outdir, "changed-other-slide"))
    wrong_style(args.golden, output_path(args.outdir, "wrong-style"))
    flattened_slide(args.source, args.slide8_image, output_path(args.outdir, "flattened"))
    duplicate_target(args.golden, output_path(args.outdir, "duplicate"))
    hardcoded_marker(args.golden, output_path(args.outdir, "hardcoded-marker"))
    moved_target(args.golden, output_path(args.outdir, "moved-target"))


if __name__ == "__main__":
    main()
