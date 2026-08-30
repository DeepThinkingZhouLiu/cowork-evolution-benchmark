#!/usr/bin/env python3
"""PPTArena Case 4 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import math
import re
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from lxml import etree
from PIL import Image, ImageChops
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "FontSizePlus2Slide8EditedA.pptx"
INPUT_NAME = "FontSizePlus2Slide8TestA.pptx"
INPUT_SHA256 = "6beba374b051d77ba11636d826535fc3e391f245daf14d7547383a02798affd6"
TASK_ID = "cowork-evo/pptarena-004-slide8-font-size-increase"
RUBRIC_VERSION = "pptarena-case-004-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-004-deterministic-v0.1"
EMU_PER_INCH = 914_400
TARGET_SLIDE_INDEX = 7
GEOMETRY_TOLERANCE = 0.012 * EMU_PER_INCH
PAGE_EDGE_TOLERANCE = 0.03 * EMU_PER_INCH

TARGET_RULES = {
    "Foranobjectoutsidethefocalpoint,concavemirrorscreateimageswhichare": [22.0],
    "Inverted(upsidedown)": [26.0, 20.0],
    "Real(formedbyintersectionofactualrays)": [26.0, 20.0],
    "Objectdistance(s)": [26.0],
    "Focallengthiscalledpositive": [22.0],
    "Canbereducedorenlarged": [26.0],
    "Imagedistance(s')": [26.0],
    "Objectheight(h)": [26.0],
    "Imageheight(h')": [26.0],
}
AUTO_HEIGHT_KEYS = {
    "Inverted(upsidedown)",
    "Real(formedbyintersectionofactualrays)",
}
RENDER_PHRASES = [
    "Foranobjectoutsidethefocalpoint,concavemirrorscreateimageswhichare",
    "Inverted",
    "upsidedown",
    "Real",
    "formedbyintersectionofactualrays",
    "Objectdistance",
    "Focallengthiscalledpositive",
    "Canbereducedorenlarged",
    "Imagedistance",
    "Objectheight",
    "Imageheight",
]
SUSPICIOUS_MARKERS = [
    "solution/golden",
    "/root/",
    "judge-result",
    "artifact_dir",
    "pptarena-case-004",
]

SPECS = {
    "R001": ("指定 PPTX 可打开，严格保留 56 页与页面尺寸", 6, "hurdle"),
    "R002": ("第 8 页九个目标文本对象的字号均精确增加 2 pt", 12, "hurdle"),
    "R003": ("目标对象位置和宽度保持，仅两个自动适应文本框可小幅增高", 8, "positive"),
    "R004": ("第 8 页全部非目标文字、例外和几何保持", 8, "hurdle"),
    "R005": ("56 页全部原有文字按原页完整保留", 8, "hurdle"),
    "R006": ("除第 8 页外的 55 页渲染保持", 8, "hurdle"),
    "R007": ("原有图片与第 8 页非文本图形保持且未压平", 5, "positive"),
    "R008": ("56 页可渲染且第 8 页目标文字完整可读、位于页面内", 8, "positive"),
    "R009": ("九个目标对象的非字号文字属性保持", 5, "hurdle"),
    "R010": ("输出目录干净且不含答案路径或未授权文字", 4, "positive"),
    "R011": ("页数、字号、例外、非目标页、原文、样式或可编辑性无严重错误", -20, "penalty"),
}


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", "", value or "")


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def locate_source() -> Path | None:
    candidates = [
        Path("/app/input_files") / INPUT_NAME,
        Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME,
    ]
    return next((path for path in candidates if path.is_file()), None)


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def txbody(shape):
    nodes = shape._element.xpath('./*[local-name()="txBody"]')
    return nodes[0] if nodes else None


def shape_has_txbody(shape) -> bool:
    return txbody(shape) is not None


def shape_text(shape) -> str:
    body = txbody(shape)
    if body is None:
        return ""
    return "\n".join(
        "".join(paragraph.xpath('.//*[local-name()="t"]/text()'))
        for paragraph in body.xpath('./*[local-name()="p"]')
    )


def geometry(shape) -> tuple[int, int, int, int]:
    return (int(shape.left), int(shape.top), int(shape.width), int(shape.height))


def rounded_geometry(shape) -> tuple[int, int, int, int]:
    return tuple(round(value / 1000) * 1000 for value in geometry(shape))


def target_shapes(slide, key: str):
    return [
        shape
        for shape in iter_shapes(slide.shapes)
        if shape_has_txbody(shape) and normalize_text(shape_text(shape)) == key
    ]


def local_attr_name(name: str) -> str:
    return etree.QName(name).localname if name.startswith("{") else name


def effective_run_size(run_node, paragraph_node, body_node) -> float | None:
    values = run_node.xpath(
        './*[local-name()="rPr" or local-name()="fldPr"]/@sz'
    )
    if not values:
        values = paragraph_node.xpath(
            './*[local-name()="pPr"]/*[local-name()="defRPr"]/@sz'
        )
    if not values:
        levels = paragraph_node.xpath('./*[local-name()="pPr"]/@lvl')
        try:
            level = int(levels[0]) + 1 if levels else 1
        except (TypeError, ValueError):
            level = 1
        values = body_node.xpath(
            f'./*[local-name()="lstStyle"]/*[local-name()="lvl{level}pPr"]/*[local-name()="defRPr"]/@sz'
        )
    if not values:
        return None
    try:
        return int(values[0]) / 100.0
    except (TypeError, ValueError):
        return None


def paragraph_sizes(shape) -> list[list[float | None]]:
    body = txbody(shape)
    if body is None:
        return []
    result: list[list[float | None]] = []
    for paragraph in body.xpath('./*[local-name()="p"]'):
        runs = []
        for run_node in paragraph.xpath(
            './*[local-name()="r" or local-name()="fld"]'
        ):
            text = "".join(run_node.xpath('./*[local-name()="t"]/text()'))
            if text:
                runs.append(effective_run_size(run_node, paragraph, body))
        if runs:
            result.append(runs)
    return result


def canonical_txbody(shape, drop_size: bool) -> bytes:
    body = txbody(shape)
    if body is None:
        return b""
    clone = copy.deepcopy(body)
    for node in clone.iter():
        for attribute in list(node.attrib):
            name = local_attr_name(attribute)
            if name in {"dirty", "smtClean"} or (drop_size and name == "sz"):
                del node.attrib[attribute]
    return etree.tostring(clone, method="c14n")


def text_shape_record(shape, drop_size: bool = False) -> tuple:
    return (
        normalize_text(shape_text(shape)),
        rounded_geometry(shape),
        hashlib.sha256(canonical_txbody(shape, drop_size)).hexdigest(),
    )


def non_target_text_counter(slide) -> Counter[tuple]:
    values: Counter[tuple] = Counter()
    for shape in iter_shapes(slide.shapes):
        if not shape_has_txbody(shape):
            continue
        if normalize_text(shape_text(shape)) in TARGET_RULES:
            continue
        values[text_shape_record(shape)] += 1
    return values


def slide_text_counter(slide) -> Counter[str]:
    values: Counter[str] = Counter()
    for shape in iter_shapes(slide.shapes):
        if not shape_has_txbody(shape):
            continue
        text = normalize_text(shape_text(shape))
        if text:
            values[text] += 1
    return values


def slide_picture_counter(slide) -> Counter[str]:
    values: Counter[str] = Counter()
    for shape in iter_shapes(slide.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            values[hashlib.sha256(shape.image.blob).hexdigest()] += 1
    return values


def non_text_shape_counter(slide) -> Counter[tuple]:
    values: Counter[tuple] = Counter()
    for shape in iter_shapes(slide.shapes):
        if shape_has_txbody(shape):
            continue
        picture_hash = ""
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            picture_hash = hashlib.sha256(shape.image.blob).hexdigest()
        values[(int(shape.shape_type), rounded_geometry(shape), picture_hash)] += 1
    return values


def package_ok(path: Path) -> tuple[bool, str]:
    try:
        with zipfile.ZipFile(path) as archive:
            broken = archive.testzip()
            names = set(archive.namelist())
            required = {
                "[Content_Types].xml",
                "ppt/presentation.xml",
                "ppt/slides/slide8.xml",
                "ppt/slides/slide56.xml",
            }
            missing = sorted(required - names)
            macros = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
        return broken is None and not missing and not macros, (
            f"损坏成员={broken}，缺失={missing}，宏={macros}"
        )
    except Exception as exc:
        return False, f"ZIP 检查失败：{exc}"


def render_presentation(path: Path, outdir: Path) -> tuple[list[Path], str, int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "lo-profile"
    profile.mkdir()
    completed = subprocess.run(
        [
            "soffice",
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(outdir),
            str(path),
        ],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=300,
        check=False,
    )
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return [], "", 0, f"转 PDF 失败：{completed.stdout[-500:]}"

    info = subprocess.run(
        ["pdfinfo", str(pdf)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=30,
        check=False,
    ).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    pages = int(match.group(1)) if match else 0

    prefix = outdir / "slide"
    raster = subprocess.run(
        ["pdftoppm", "-jpeg", "-r", "90", str(pdf), str(prefix)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=360,
        check=False,
    )
    images = sorted(outdir.glob("slide-*.jpg")) if raster.returncode == 0 else []

    page8_text = outdir / "page8.txt"
    extracted = subprocess.run(
        ["pdftotext", "-f", "8", "-l", "8", str(pdf), str(page8_text)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=30,
        check=False,
    )
    text = (
        page8_text.read_text(encoding="utf-8", errors="replace")
        if extracted.returncode == 0 and page8_text.is_file()
        else ""
    )
    evidence = f"页数={pages}，渲染图={len(images)}，pdftotext={extracted.returncode}"
    return images, text, pages, evidence


def image_similarity(first: Path, second: Path) -> float:
    with Image.open(first) as left_image, Image.open(second) as right_image:
        left = left_image.convert("RGB")
        right = right_image.convert("RGB")
        if left.size != right.size:
            return 0.0
        difference = ImageChops.difference(left, right)
        histogram = difference.histogram()
        squared = sum(
            count * ((index % 256) ** 2)
            for index, count in enumerate(histogram)
        )
        rms = math.sqrt(squared / (left.width * left.height * 3))
        return max(0.0, 1.0 - rms / 255.0)


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {
        "criterion_id": identifier,
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
        "raw": {
            "description": description,
            "weight": weight,
            "type": kind,
            "passed": bool(passed),
        },
    }


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [
        item["criterion_id"]
        for item in criteria
        if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]
    ]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(
        item["raw"]["weight"]
        for item in criteria
        if item["raw"]["weight"] > 0 and item["raw"]["passed"]
    )
    penalty = sum(
        abs(item["raw"]["weight"])
        for item in criteria
        if item["raw"]["weight"] < 0 and not item["raw"]["passed"]
    )
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [sha256(path)] if path.is_file() else [],
        "criterion_results": criteria,
        "reward": reward,
        "passed": reward >= 0.8 and not hurdles,
        "hurdle_failures": hurdles,
    }


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        criteria = [result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配") for identifier in SPECS]
        return make_result(path, criteria)

    try:
        candidate = Presentation(path)
        source = Presentation(source_path)
    except Exception as exc:
        criteria = [result_item(identifier, False, f"PPTX 打开失败：{exc}") for identifier in SPECS]
        return make_result(path, criteria)

    criteria: list[dict] = []
    package_passed, package_evidence = package_ok(path)
    size_passed = (
        len(candidate.slides) == len(source.slides) == 56
        and candidate.slide_width == source.slide_width
        and candidate.slide_height == source.slide_height
    )
    criteria.append(
        result_item(
            "R001",
            package_passed and size_passed,
            f"{package_evidence}，页数={len(candidate.slides)}，页面={candidate.slide_width/EMU_PER_INCH:.4f}×{candidate.slide_height/EMU_PER_INCH:.4f}英寸",
        )
    )

    target_size_failures: list[str] = []
    target_geometry_failures: list[str] = []
    target_style_failures: list[str] = []
    target_candidates = {}
    if len(candidate.slides) > TARGET_SLIDE_INDEX:
        source_slide = source.slides[TARGET_SLIDE_INDEX]
        candidate_slide = candidate.slides[TARGET_SLIDE_INDEX]
        for key, expected_by_paragraph in TARGET_RULES.items():
            source_matches = target_shapes(source_slide, key)
            candidate_matches = target_shapes(candidate_slide, key)
            if len(source_matches) != 1 or len(candidate_matches) != 1:
                target_size_failures.append(
                    f"{key}:源={len(source_matches)},候选={len(candidate_matches)}"
                )
                target_geometry_failures.append(key)
                target_style_failures.append(key)
                continue
            source_shape = source_matches[0]
            candidate_shape = candidate_matches[0]
            target_candidates[key] = candidate_shape

            actual_by_paragraph = paragraph_sizes(candidate_shape)
            size_ok = len(actual_by_paragraph) == len(expected_by_paragraph)
            if size_ok:
                for actual_runs, expected in zip(actual_by_paragraph, expected_by_paragraph):
                    if not actual_runs or any(
                        actual is None or abs(actual - expected) > 0.05
                        for actual in actual_runs
                    ):
                        size_ok = False
                        break
            if not size_ok:
                target_size_failures.append(
                    f"{key}:实际={actual_by_paragraph},期望={expected_by_paragraph}"
                )

            sx, sy, sw, sh = geometry(source_shape)
            cx, cy, cw, ch = geometry(candidate_shape)
            fixed_ok = all(
                abs(actual - expected) <= GEOMETRY_TOLERANCE
                for actual, expected in [(cx, sx), (cy, sy), (cw, sw)]
            )
            if key in AUTO_HEIGHT_KEYS:
                height_ok = (
                    ch >= sh - GEOMETRY_TOLERANCE
                    and ch <= sh + int(0.22 * EMU_PER_INCH)
                    and cy + ch <= candidate.slide_height + PAGE_EDGE_TOLERANCE
                )
            else:
                height_ok = abs(ch - sh) <= GEOMETRY_TOLERANCE
            if not fixed_ok or not height_ok:
                target_geometry_failures.append(
                    f"{key}:源={(sx, sy, sw, sh)},候选={(cx, cy, cw, ch)}"
                )

            if canonical_txbody(source_shape, drop_size=True) != canonical_txbody(
                candidate_shape, drop_size=True
            ):
                target_style_failures.append(key)
    else:
        target_size_failures.append("缺失第 8 页")
        target_geometry_failures.append("缺失第 8 页")
        target_style_failures.append("缺失第 8 页")

    criteria.append(
        result_item(
            "R002",
            not target_size_failures and len(target_candidates) == len(TARGET_RULES),
            f"目标对象={len(target_candidates)}/9，字号失败={target_size_failures}",
        )
    )
    criteria.append(
        result_item(
            "R003",
            not target_geometry_failures and len(target_candidates) == len(TARGET_RULES),
            f"几何失败={target_geometry_failures}",
        )
    )

    non_target_passed = False
    non_target_evidence = "缺失第 8 页"
    if len(candidate.slides) > TARGET_SLIDE_INDEX:
        source_non_target = non_target_text_counter(source.slides[TARGET_SLIDE_INDEX])
        candidate_non_target = non_target_text_counter(candidate.slides[TARGET_SLIDE_INDEX])
        missing = list((source_non_target - candidate_non_target).elements())
        extra = list((candidate_non_target - source_non_target).elements())
        non_target_passed = not missing and not extra
        non_target_evidence = f"缺失对象={len(missing)}，新增或变化对象={len(extra)}"
    criteria.append(result_item("R004", non_target_passed, non_target_evidence))

    source_texts = [slide_text_counter(slide) for slide in source.slides]
    candidate_texts = [slide_text_counter(slide) for slide in candidate.slides]
    if len(source_texts) == len(candidate_texts):
        text_mismatches = [
            index + 1
            for index, (actual, expected) in enumerate(zip(candidate_texts, source_texts))
            if actual != expected
        ]
    else:
        text_mismatches = list(range(1, max(len(source_texts), len(candidate_texts)) + 1))
    text_passed = not text_mismatches
    criteria.append(result_item("R005", text_passed, f"文字不一致页={text_mismatches}"))

    source_pictures = [slide_picture_counter(slide) for slide in source.slides]
    candidate_pictures = [slide_picture_counter(slide) for slide in candidate.slides]
    if len(source_pictures) == len(candidate_pictures):
        picture_mismatches = [
            index + 1
            for index, (actual, expected) in enumerate(zip(candidate_pictures, source_pictures))
            if actual != expected
        ]
    else:
        picture_mismatches = list(range(1, max(len(source_pictures), len(candidate_pictures)) + 1))

    full_slide_new = []
    for slide_index, slide in enumerate(candidate.slides, start=1):
        for shape in iter_shapes(slide.shapes):
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            coverage = (shape.width * shape.height) / (candidate.slide_width * candidate.slide_height)
            digest = hashlib.sha256(shape.image.blob).hexdigest()
            source_counter = source_pictures[slide_index - 1] if slide_index <= len(source_pictures) else Counter()
            if coverage >= 0.9 and source_counter[digest] == 0:
                full_slide_new.append(slide_index)

    non_text_mismatch = True
    if len(candidate.slides) > TARGET_SLIDE_INDEX:
        non_text_mismatch = (
            non_text_shape_counter(candidate.slides[TARGET_SLIDE_INDEX])
            != non_text_shape_counter(source.slides[TARGET_SLIDE_INDEX])
        )
    native_passed = not picture_mismatches and not full_slide_new and not non_text_mismatch
    criteria.append(
        result_item(
            "R007",
            native_passed,
            f"图片不一致页={picture_mismatches}，新增整页图片页={full_slide_new}，第8页非文本图形变化={non_text_mismatch}",
        )
    )

    style_passed = not target_style_failures and len(target_candidates) == len(TARGET_RULES)
    criteria.append(result_item("R009", style_passed, f"非字号样式失败={target_style_failures}"))

    artifact_files = sorted(
        str(item.relative_to(artifact_dir))
        for item in artifact_dir.rglob("*")
        if item.is_file()
    )
    all_text = "\n".join(shape_text(shape) for slide in candidate.slides for shape in iter_shapes(slide.shapes))
    lowered = all_text.lower()
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker.lower() in lowered]
    hygiene_passed = artifact_files == [OUTPUT_NAME] and not markers
    criteria.append(
        result_item(
            "R010",
            hygiene_passed,
            f"输出文件={artifact_files}，可疑标记={markers}",
        )
    )

    core_ready = all(
        next(item for item in criteria if item["criterion_id"] == identifier)["raw"]["passed"]
        for identifier in ["R001", "R002", "R004", "R005", "R009"]
    )
    non_target_render_passed = False
    render_passed = False
    non_target_min = 0.0
    render_evidence = "前置结构或内容检查失败，未执行渲染"
    missing_render_phrases = list(RENDER_PHRASES)
    target_contained = False

    if core_ready:
        with tempfile.TemporaryDirectory(prefix="pptarena004-judge-") as temp_dir:
            root = Path(temp_dir)
            source_images, _, source_pages, source_note = render_presentation(source_path, root / "source")
            candidate_images, page8_text, candidate_pages, candidate_note = render_presentation(path, root / "candidate")
            render_evidence = f"源[{source_note}]；候选[{candidate_note}]"

            non_target_scores: list[float] = []
            if len(source_images) == 56 and len(candidate_images) == 56:
                for index in range(56):
                    if index != TARGET_SLIDE_INDEX:
                        non_target_scores.append(
                            image_similarity(source_images[index], candidate_images[index])
                        )
            non_target_min = min(non_target_scores) if non_target_scores else 0.0
            non_target_render_passed = (
                len(non_target_scores) == 55 and non_target_min >= 0.999
            )

            normalized_render = normalize_text(page8_text)
            missing_render_phrases = [
                phrase for phrase in RENDER_PHRASES if phrase not in normalized_render
            ]
            target_contained = len(target_candidates) == len(TARGET_RULES) and all(
                shape.left >= -PAGE_EDGE_TOLERANCE
                and shape.top >= -PAGE_EDGE_TOLERANCE
                and shape.left + shape.width <= candidate.slide_width + PAGE_EDGE_TOLERANCE
                and shape.top + shape.height <= candidate.slide_height + PAGE_EDGE_TOLERANCE
                for shape in target_candidates.values()
            )
            render_passed = (
                source_pages == 56
                and candidate_pages == 56
                and len(candidate_images) == 56
                and not missing_render_phrases
                and target_contained
            )

    criteria.append(
        result_item(
            "R006",
            non_target_render_passed,
            f"55 页最低像素相似度={non_target_min:.6f}；{render_evidence}",
        )
    )
    criteria.append(
        result_item(
            "R008",
            render_passed,
            f"缺失可提取短语={missing_render_phrases}，目标框均在页内={target_contained}；{render_evidence}",
        )
    )

    severe = (
        not core_ready
        or not non_target_render_passed
        or bool(full_slide_new)
        or bool(markers)
        or not native_passed
    )
    criteria.append(result_item("R011", not severe, f"严重错误={severe}"))

    order = {identifier: index for index, identifier in enumerate(SPECS)}
    criteria.sort(key=lambda item: order[item["criterion_id"]])
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()

    output = args.artifact_dir / OUTPUT_NAME
    if output.is_file():
        evaluation = evaluate(output, args.artifact_dir)
    else:
        evaluation = {
            "task_id": TASK_ID,
            "rubric_version": RUBRIC_VERSION,
            "judge_version": JUDGE_VERSION,
            "artifact_sha256": [],
            "criterion_results": [],
            "reward": 0.0,
            "passed": False,
            "hurdle_failures": ["missing_output"],
        }

    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
