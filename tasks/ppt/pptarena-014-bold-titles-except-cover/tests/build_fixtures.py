#!/usr/bin/env python3
"""只为 PPTArena Case 14 构造可重放的 Judge 负例和中等例。"""

from __future__ import annotations

import argparse
import re
import shutil
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches


OUTPUT_NAME = "Bolded_BoldTitleTestA.pptx"


def output_path(root: Path, case: str) -> Path:
    target = root / case / OUTPUT_NAME
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def set_title_bold(
    presentation: Presentation,
    slide_numbers: list[int],
    value: bool | None,
) -> None:
    for slide_number in slide_numbers:
        title = presentation.slides[slide_number - 1].shapes.title
        if title is None:
            raise ValueError(f"第 {slide_number} 页缺少主标题")
        changed = 0
        for paragraph in title.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.text.strip():
                    run.font.bold = value
                    changed += 1
        if changed == 0:
            raise ValueError(f"第 {slide_number} 页标题没有非空文字运行")


def skip_one_title(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    set_title_bold(presentation, [17], False)
    presentation.save(output)


def skip_hidden_titles(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    set_title_bold(presentation, [8, 9, 23, 31], None)
    presentation.save(output)


def bold_cover(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    set_title_bold(presentation, [1], True)
    presentation.save(output)


def bold_everything(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    for slide in presentation.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip():
                        run.font.bold = True
    presentation.save(output)


def wrong_title_style(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    title = presentation.slides[11].shapes.title
    for paragraph in title.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.text.strip():
                run.font.color.rgb = RGBColor(255, 0, 0)
                presentation.save(output)
                return
    raise ValueError("第 12 页标题没有可变异文字运行")


def changed_body(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    slide = presentation.slides[19]
    title_element = slide.shapes.title._element
    for shape in slide.shapes:
        if (
            shape._element is title_element
            or not getattr(shape, "has_text_frame", False)
        ):
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.text.strip():
                    run.text += " X"
                    presentation.save(output)
                    return
    raise ValueError("第 20 页没有可变异正文")


def changed_hidden_status(golden: Path, output: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="pptarena014-hidden-") as temporary:
        temporary_path = Path(temporary) / OUTPUT_NAME
        with zipfile.ZipFile(golden) as reader, zipfile.ZipFile(
            temporary_path, "w", compression=zipfile.ZIP_DEFLATED
        ) as writer:
            for info in reader.infolist():
                data = reader.read(info.filename)
                if info.filename == "ppt/slides/slide8.xml":
                    root = etree.fromstring(data)
                    root.attrib.pop("show", None)
                    data = etree.tostring(
                        root,
                        xml_declaration=True,
                        encoding="UTF-8",
                        standalone=True,
                    )
                writer.writestr(info, data)
        shutil.copy2(temporary_path, output)


def flattened_slide(
    golden: Path,
    slide_image: Path,
    output: Path,
) -> None:
    presentation = Presentation(golden)
    slide = presentation.slides[9]
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    slide.shapes.add_picture(
        str(slide_image),
        0,
        0,
        width=presentation.slide_width,
        height=presentation.slide_height,
    )
    presentation.save(output)


def hardcoded_marker(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    box = presentation.slides[0].shapes.add_textbox(
        Inches(0.1), Inches(7.6), Inches(2.5), Inches(0.2)
    )
    box.text = "solution/Golden"
    presentation.save(output)


def moved_title(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    title = presentation.slides[16].shapes.title
    title.top += Inches(0.08)
    presentation.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--slide10-image", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()

    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))
    skip_one_title(args.golden, output_path(args.outdir, "skip-one-title"))
    skip_hidden_titles(
        args.golden, output_path(args.outdir, "skip-hidden-titles")
    )
    bold_cover(args.golden, output_path(args.outdir, "bold-cover"))
    bold_everything(args.golden, output_path(args.outdir, "bold-everything"))
    wrong_title_style(
        args.golden, output_path(args.outdir, "wrong-title-style")
    )
    changed_body(args.golden, output_path(args.outdir, "changed-body"))
    changed_hidden_status(
        args.golden, output_path(args.outdir, "changed-hidden-status")
    )
    flattened_slide(
        args.golden,
        args.slide10_image,
        output_path(args.outdir, "flattened"),
    )
    hardcoded_marker(
        args.golden, output_path(args.outdir, "hardcoded-marker")
    )
    moved_title(args.golden, output_path(args.outdir, "moved-title"))


if __name__ == "__main__":
    main()
