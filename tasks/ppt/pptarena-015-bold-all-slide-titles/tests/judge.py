#!/usr/bin/env python3
"""PPTArena Case 15 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import math
import re
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from lxml import etree
from PIL import Image, ImageChops, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "Bolded_BoldTitleTestB.pptx"
INPUT_NAME = "BoldTitleTestB_cse1300-ai.pptx"
INPUT_SHA256 = "02687b50c86639c32501eefdf04cfb45e1687f45f9502c7896f4bf46f9e892b0"
TASK_ID = "cowork-evo/pptarena-015-bold-all-slide-titles"
RUBRIC_VERSION = "pptarena-case-015-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-015-deterministic-v0.1"
EMU_PER_INCH = 914_400
GEOMETRY_TOLERANCE = int(0.012 * EMU_PER_INCH)
PAGE_EDGE_TOLERANCE = int(0.03 * EMU_PER_INCH)
SUSPICIOUS_MARKERS = [
    "solution/golden",
    "/root/",
    "judge-result",
    "artifact_dir",
    "pptarena-case-015",
]

SPECS = {
    "R001": ("指定 PPTX 可打开，严格保留 11 页与页面尺寸", 6, "hurdle"),
    "R002": ("第 1–11 页全部主标题均为粗体", 16, "hurdle"),
    "R003": ("第 1 页课程副标题和其他非主标题元素保持", 8, "hurdle"),
    "R004": ("11 页标题除粗体外的内容和样式保持", 9, "hurdle"),
    "R005": ("全部标题的占位符语义和几何保持", 7, "positive"),
    "R006": ("11 页原文与讲者备注按页完整保持", 9, "hurdle"),
    "R007": ("全部非标题文字对象的样式、几何和层级保持", 8, "hurdle"),
    "R008": ("图片和非文本图形保持且未压平", 6, "hurdle"),
    "R009": ("11 页可渲染且标题外视觉保持", 9, "positive"),
    "R010": ("输出目录干净且不含答案路径或未授权文字", 4, "positive"),
    "R011": ("页数、目标标题、封面副标题、正文或可编辑性无严重错误", -20, "penalty"),
}


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", "", value or "")


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def locate_source() -> Path | None:
    candidates = [
        Path("/app/input_files") / INPUT_NAME,
        Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME,
    ]
    return next((path for path in candidates if path.is_file()), None)


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def title_shape(slide):
    try:
        return slide.shapes.title
    except Exception:
        return None


def txbody(shape):
    if shape is None:
        return None
    nodes = shape._element.xpath('./*[local-name()="txBody"]')
    return nodes[0] if nodes else None


def shape_has_txbody(shape) -> bool:
    return txbody(shape) is not None


def shape_text(shape) -> str:
    body = txbody(shape)
    if body is None:
        return ""
    return "\n".join(
        "".join(paragraph.xpath('.//*[local-name()="t"]/text()'))
        for paragraph in body.xpath('./*[local-name()="p"]')
    )


def geometry(shape) -> tuple[int, int, int, int]:
    return (int(shape.left), int(shape.top), int(shape.width), int(shape.height))


def rounded_geometry(shape) -> tuple[int, int, int, int]:
    return tuple(round(value / 1000) * 1000 for value in geometry(shape))


def local_attr_name(name: str) -> str:
    return etree.QName(name).localname if name.startswith("{") else name


def canonical_txbody(shape, drop_bold: bool) -> bytes:
    body = txbody(shape)
    if body is None:
        return b""
    clone = copy.deepcopy(body)
    for node in clone.iter():
        for attribute in list(node.attrib):
            name = local_attr_name(attribute)
            if name in {"dirty", "smtClean"} or (drop_bold and name == "b"):
                del node.attrib[attribute]
    return etree.tostring(clone, method="c14n")


def xml_bool(value: str | None) -> bool | None:
    if value is None:
        return None
    lowered = value.strip().lower()
    if lowered in {"1", "true", "on"}:
        return True
    if lowered in {"0", "false", "off"}:
        return False
    return None


def effective_run_bold(run_node, paragraph_node, body_node) -> bool | None:
    values = run_node.xpath(
        './*[local-name()="rPr" or local-name()="fldPr"]/@b'
    )
    if not values:
        values = paragraph_node.xpath(
            './*[local-name()="pPr"]/*[local-name()="defRPr"]/@b'
        )
    if not values:
        levels = paragraph_node.xpath('./*[local-name()="pPr"]/@lvl')
        try:
            level = int(levels[0]) + 1 if levels else 1
        except (TypeError, ValueError):
            level = 1
        values = body_node.xpath(
            f'./*[local-name()="lstStyle"]/*[local-name()="lvl{level}pPr"]/*[local-name()="defRPr"]/@b'
        )
    return xml_bool(values[0]) if values else None


def title_bold_values(shape) -> list[bool | None]:
    body = txbody(shape)
    if body is None:
        return []
    values: list[bool | None] = []
    for paragraph in body.xpath('./*[local-name()="p"]'):
        for run_node in paragraph.xpath(
            './*[local-name()="r" or local-name()="fld"]'
        ):
            text = "".join(run_node.xpath('./*[local-name()="t"]/text()'))
            if text.strip():
                values.append(effective_run_bold(run_node, paragraph, body))
    return values


def placeholder_type(shape) -> str:
    if shape is None or not getattr(shape, "is_placeholder", False):
        return ""
    return str(shape.placeholder_format.type)


def text_shape_record(shape) -> tuple:
    return (
        int(shape.shape_type),
        placeholder_type(shape),
        normalize_text(shape_text(shape)),
        rounded_geometry(shape),
        hashlib.sha256(canonical_txbody(shape, drop_bold=False)).hexdigest(),
    )


def non_title_text_records(slide) -> list[tuple]:
    title = title_shape(slide)
    title_element = title._element if title is not None else None
    records: list[tuple] = []
    for shape in iter_shapes(slide.shapes):
        if shape._element is title_element or not shape_has_txbody(shape):
            continue
        records.append(text_shape_record(shape))
    return records


def slide_text_counter(slide) -> Counter[str]:
    values: Counter[str] = Counter()
    for shape in iter_shapes(slide.shapes):
        if not shape_has_txbody(shape):
            continue
        text = normalize_text(shape_text(shape))
        if text:
            values[text] += 1
    return values


def slide_notes(slide) -> tuple[str, ...]:
    try:
        return tuple(
            shape.text
            for shape in slide.notes_slide.shapes
            if getattr(shape, "has_text_frame", False)
        )
    except Exception:
        return ()


def slide_picture_counter(slide) -> Counter[tuple]:
    values: Counter[tuple] = Counter()
    for shape in iter_shapes(slide.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            values[
                (
                    hashlib.sha256(shape.image.blob).hexdigest(),
                    rounded_geometry(shape),
                )
            ] += 1
    return values


def non_text_shape_records(slide) -> list[tuple]:
    values: list[tuple] = []
    for shape in iter_shapes(slide.shapes):
        if shape_has_txbody(shape):
            continue
        picture_hash = ""
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            picture_hash = hashlib.sha256(shape.image.blob).hexdigest()
        values.append(
            (int(shape.shape_type), rounded_geometry(shape), picture_hash)
        )
    return values


def media_hashes(path: Path) -> Counter[str]:
    with zipfile.ZipFile(path) as archive:
        return Counter(
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("ppt/media/")
        )


def package_ok(path: Path) -> tuple[bool, str]:
    try:
        with zipfile.ZipFile(path) as archive:
            broken = archive.testzip()
            names = set(archive.namelist())
            required = {
                "[Content_Types].xml",
                "ppt/presentation.xml",
                "ppt/slides/slide1.xml",
                "ppt/slides/slide11.xml",
            }
            missing = sorted(required - names)
            macros = sorted(
                name for name in names if name.lower().endswith("vbaproject.bin")
            )
        return broken is None and not missing and not macros, (
            f"损坏成员={broken}，缺失={missing}，宏={macros}"
        )
    except Exception as exc:
        return False, f"ZIP 检查失败：{exc}"


def render_presentation(
    path: Path, outdir: Path
) -> tuple[list[Path], int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "lo-profile"
    profile.mkdir()
    completed = subprocess.run(
        [
            "soffice",
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(outdir),
            str(path),
        ],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return [], 0, f"转 PDF 失败：{completed.stdout[-500:]}"
    info = subprocess.run(
        ["pdfinfo", str(pdf)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=30,
        check=False,
    ).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    pages = int(match.group(1)) if match else 0
    prefix = outdir / "slide"
    raster = subprocess.run(
        ["pdftoppm", "-jpeg", "-r", "90", str(pdf), str(prefix)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    images = sorted(outdir.glob("slide-*.jpg")) if raster.returncode == 0 else []
    return images, pages, (
        f"页数={pages}，渲染图={len(images)}，"
        f"soffice={completed.returncode}，pdftoppm={raster.returncode}"
    )


def emu_rect_to_pixels(
    shape,
    image_size: tuple[int, int],
    slide_size: tuple[int, int],
    padding: int = 48,
) -> tuple[int, int, int, int]:
    image_width, image_height = image_size
    slide_width, slide_height = slide_size
    left, top, width, height = geometry(shape)
    return (
        max(0, math.floor(left * image_width / slide_width) - padding),
        max(0, math.floor(top * image_height / slide_height) - padding),
        min(
            image_width,
            math.ceil((left + width) * image_width / slide_width) + padding,
        ),
        min(
            image_height,
            math.ceil((top + height) * image_height / slide_height) + padding,
        ),
    )


def page_overflow(
    shape, slide_size: tuple[int, int]
) -> tuple[int, int, int, int]:
    slide_width, slide_height = slide_size
    left, top, width, height = geometry(shape)
    return (
        max(0, -left),
        max(0, -top),
        max(0, left + width - slide_width),
        max(0, top + height - slide_height),
    )


def image_similarity(
    first: Path,
    second: Path,
    ignored_rectangles: list[tuple[int, int, int, int]],
) -> float:
    with Image.open(first) as left_image, Image.open(second) as right_image:
        left = left_image.convert("RGB")
        right = right_image.convert("RGB")
        if left.size != right.size:
            return 0.0
        difference = ImageChops.difference(left, right)
        mask = Image.new("L", left.size, 255)
        draw = ImageDraw.Draw(mask)
        for rectangle in ignored_rectangles:
            draw.rectangle(rectangle, fill=0)
        rgb_mask = Image.merge("RGB", (mask, mask, mask))
        masked = ImageChops.multiply(difference, rgb_mask)
        histogram = masked.histogram()
        squared = sum(
            count * ((index % 256) ** 2)
            for index, count in enumerate(histogram)
        )
        unmasked_pixels = mask.histogram()[255]
        if unmasked_pixels <= 0:
            return 0.0
        rms = math.sqrt(squared / (unmasked_pixels * 3))
        return max(0.0, 1.0 - rms / 255.0)


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {
        "criterion_id": identifier,
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
        "raw": {
            "description": description,
            "weight": weight,
            "type": kind,
            "passed": bool(passed),
        },
    }


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [
        item["criterion_id"]
        for item in criteria
        if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]
    ]
    total = sum(
        item["raw"]["weight"]
        for item in criteria
        if item["raw"]["weight"] > 0
    )
    earned = sum(
        item["raw"]["weight"]
        for item in criteria
        if item["raw"]["weight"] > 0 and item["raw"]["passed"]
    )
    penalty = sum(
        abs(item["raw"]["weight"])
        for item in criteria
        if item["raw"]["weight"] < 0 and not item["raw"]["passed"]
    )
    reward = 0.0 if hurdles else max(
        0.0, min(1.0, (earned - penalty) / total)
    )
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [sha256(path)] if path.is_file() else [],
        "criterion_results": criteria,
        "reward": reward,
        "passed": reward >= 0.8 and not hurdles,
        "hurdle_failures": hurdles,
    }


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        criteria = [
            result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配")
            for identifier in SPECS
        ]
        return make_result(path, criteria)
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        criteria = [
            result_item(identifier, False, f"PPTX 打开失败：{exc}")
            for identifier in SPECS
        ]
        return make_result(path, criteria)

    criteria: list[dict] = []
    package_passed, package_evidence = package_ok(path)
    size_passed = (
        len(candidate.slides) == len(source.slides) == 11
        and candidate.slide_width == source.slide_width
        and candidate.slide_height == source.slide_height
    )
    criteria.append(
        result_item(
            "R001",
            package_passed and size_passed,
            f"{package_evidence}，页数={len(candidate.slides)}，"
            f"页面={candidate.slide_width / EMU_PER_INCH:.4f}×"
            f"{candidate.slide_height / EMU_PER_INCH:.4f}英寸",
        )
    )

    bold_failures: list[str] = []
    style_failures: list[int] = []
    geometry_failures: list[int] = []
    candidate_titles = []
    if len(candidate.slides) == len(source.slides) == 11:
        for slide_number, (source_slide, candidate_slide) in enumerate(
            zip(source.slides, candidate.slides), start=1
        ):
            source_title = title_shape(source_slide)
            candidate_title = title_shape(candidate_slide)
            candidate_titles.append(candidate_title)
            if source_title is None or candidate_title is None:
                bold_failures.append(
                    f"第{slide_number}页缺失主标题占位符"
                )
                style_failures.append(slide_number)
                geometry_failures.append(slide_number)
                continue
            bold_values = title_bold_values(candidate_title)
            if not bold_values or any(value is not True for value in bold_values):
                bold_failures.append(
                    f"第{slide_number}页粗体值={bold_values}"
                )
            if (
                canonical_txbody(candidate_title, drop_bold=True)
                != canonical_txbody(source_title, drop_bold=True)
            ):
                style_failures.append(slide_number)
            geometry_ok = all(
                abs(actual - expected) <= GEOMETRY_TOLERANCE
                for actual, expected in zip(
                    geometry(candidate_title), geometry(source_title)
                )
            )
            placeholder_ok = (
                placeholder_type(candidate_title)
                == placeholder_type(source_title)
                and bool(placeholder_type(source_title))
            )
            if not geometry_ok or not placeholder_ok:
                geometry_failures.append(slide_number)
    else:
        bold_failures.append("页数不等于 11，无法逐页检查标题")
        style_failures = list(range(1, 12))
        geometry_failures = list(range(1, 12))

    criteria.append(
        result_item("R002", not bold_failures, f"标题粗体失败={bold_failures}")
    )

    cover_non_title_passed = False
    cover_evidence = "缺失第 1 页"
    if len(candidate.slides) >= 1 and len(source.slides) >= 1:
        source_cover = non_title_text_records(source.slides[0])
        candidate_cover = non_title_text_records(candidate.slides[0])
        cover_non_title_passed = source_cover == candidate_cover
        cover_evidence = (
            f"输入非主标题对象={len(source_cover)}，"
            f"候选非主标题对象={len(candidate_cover)}，"
            f"完全一致={cover_non_title_passed}"
        )
    criteria.append(
        result_item("R003", cover_non_title_passed, cover_evidence)
    )
    criteria.append(
        result_item(
            "R004",
            not style_failures,
            f"标题非粗体属性变化页={sorted(set(style_failures))}",
        )
    )
    criteria.append(
        result_item(
            "R005",
            not geometry_failures,
            f"标题几何或占位符变化页={sorted(set(geometry_failures))}",
        )
    )

    source_texts = [slide_text_counter(slide) for slide in source.slides]
    candidate_texts = [slide_text_counter(slide) for slide in candidate.slides]
    text_mismatches = (
        [
            index + 1
            for index, (actual, expected) in enumerate(
                zip(candidate_texts, source_texts)
            )
            if actual != expected
        ]
        if len(candidate_texts) == len(source_texts)
        else list(range(1, 12))
    )
    source_notes = [slide_notes(slide) for slide in source.slides]
    candidate_notes = [slide_notes(slide) for slide in candidate.slides]
    notes_mismatches = (
        [
            index + 1
            for index, (actual, expected) in enumerate(
                zip(candidate_notes, source_notes)
            )
            if actual != expected
        ]
        if len(candidate_notes) == len(source_notes)
        else list(range(1, 12))
    )
    content_passed = not text_mismatches and not notes_mismatches
    criteria.append(
        result_item(
            "R006",
            content_passed,
            f"文字不一致页={text_mismatches}，备注不一致页={notes_mismatches}",
        )
    )

    non_title_mismatches = (
        [
            index + 1
            for index, (candidate_slide, source_slide) in enumerate(
                zip(candidate.slides, source.slides)
            )
            if non_title_text_records(candidate_slide)
            != non_title_text_records(source_slide)
        ]
        if len(candidate.slides) == len(source.slides)
        else list(range(1, 12))
    )
    criteria.append(
        result_item(
            "R007",
            not non_title_mismatches,
            f"非标题文字对象变化页={non_title_mismatches}",
        )
    )

    picture_mismatches = (
        [
            index + 1
            for index, (candidate_slide, source_slide) in enumerate(
                zip(candidate.slides, source.slides)
            )
            if slide_picture_counter(candidate_slide)
            != slide_picture_counter(source_slide)
        ]
        if len(candidate.slides) == len(source.slides)
        else list(range(1, 12))
    )
    non_text_mismatches = (
        [
            index + 1
            for index, (candidate_slide, source_slide) in enumerate(
                zip(candidate.slides, source.slides)
            )
            if non_text_shape_records(candidate_slide)
            != non_text_shape_records(source_slide)
        ]
        if len(candidate.slides) == len(source.slides)
        else list(range(1, 12))
    )
    full_slide_new: list[int] = []
    for slide_index, slide in enumerate(candidate.slides, start=1):
        source_pictures = (
            slide_picture_counter(source.slides[slide_index - 1])
            if slide_index <= len(source.slides)
            else Counter()
        )
        for shape in iter_shapes(slide.shapes):
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            coverage = (
                shape.width * shape.height
                / (candidate.slide_width * candidate.slide_height)
            )
            key = (
                hashlib.sha256(shape.image.blob).hexdigest(),
                rounded_geometry(shape),
            )
            if coverage >= 0.9 and source_pictures[key] == 0:
                full_slide_new.append(slide_index)
    try:
        media_equal = media_hashes(path) == media_hashes(source_path)
    except Exception:
        media_equal = False
    native_passed = (
        not picture_mismatches
        and not non_text_mismatches
        and not full_slide_new
        and media_equal
    )
    criteria.append(
        result_item(
            "R008",
            native_passed,
            f"图片变化页={picture_mismatches}，"
            f"非文本对象变化页={non_text_mismatches}，"
            f"新增整页图片页={full_slide_new}，媒体哈希集合一致={media_equal}",
        )
    )

    artifact_files = sorted(
        str(item.relative_to(artifact_dir))
        for item in artifact_dir.rglob("*")
        if item.is_file()
    )
    all_text = "\n".join(
        shape_text(shape)
        for slide in candidate.slides
        for shape in iter_shapes(slide.shapes)
    )
    all_notes = "\n".join(
        text for slide in candidate.slides for text in slide_notes(slide)
    )
    lowered = f"{all_text}\n{all_notes}".lower()
    markers = [
        marker for marker in SUSPICIOUS_MARKERS if marker.lower() in lowered
    ]
    hygiene_passed = artifact_files == [OUTPUT_NAME] and not markers
    criteria.append(
        result_item(
            "R010",
            hygiene_passed,
            f"输出文件={artifact_files}，可疑标记={markers}",
        )
    )

    core_ready = all(
        next(
            item
            for item in criteria
            if item["criterion_id"] == identifier
        )["raw"]["passed"]
        for identifier in [
            "R001",
            "R002",
            "R003",
            "R004",
            "R006",
            "R007",
            "R008",
        ]
    )
    render_passed = False
    render_min = 0.0
    title_bounds_failures: list[int] = []
    render_evidence = "前置结构或内容检查失败，未执行渲染"

    if core_ready:
        with tempfile.TemporaryDirectory(
            prefix="pptarena015-judge-"
        ) as temporary:
            root = Path(temporary)
            source_images, source_pages, source_note = render_presentation(
                source_path, root / "source"
            )
            candidate_images, candidate_pages, candidate_note = render_presentation(
                path, root / "candidate"
            )
            render_evidence = f"源[{source_note}]；候选[{candidate_note}]"
            scores: list[float] = []
            if len(source_images) == len(candidate_images) == 11:
                for index, (source_image, candidate_image) in enumerate(
                    zip(source_images, candidate_images)
                ):
                    with Image.open(source_image) as image:
                        image_size = image.size
                    source_title = title_shape(source.slides[index])
                    candidate_title = title_shape(candidate.slides[index])
                    ignored: list[tuple[int, int, int, int]] = []
                    if source_title is not None:
                        ignored.append(
                            emu_rect_to_pixels(
                                source_title,
                                image_size,
                                (source.slide_width, source.slide_height),
                            )
                        )
                    if candidate_title is not None:
                        ignored.append(
                            emu_rect_to_pixels(
                                candidate_title,
                                image_size,
                                (
                                    candidate.slide_width,
                                    candidate.slide_height,
                                ),
                            )
                        )
                    scores.append(
                        image_similarity(
                            source_image, candidate_image, ignored
                        )
                    )
            render_min = min(scores) if scores else 0.0

            for slide_number, shape in enumerate(candidate_titles, start=1):
                if shape is None:
                    title_bounds_failures.append(slide_number)
                    continue
                source_shape = title_shape(source.slides[slide_number - 1])
                if source_shape is None:
                    title_bounds_failures.append(slide_number)
                    continue
                source_overflow = page_overflow(
                    source_shape, (source.slide_width, source.slide_height)
                )
                candidate_overflow = page_overflow(
                    shape, (candidate.slide_width, candidate.slide_height)
                )
                if any(
                    actual > expected + PAGE_EDGE_TOLERANCE
                    for actual, expected in zip(
                        candidate_overflow, source_overflow
                    )
                ):
                    title_bounds_failures.append(slide_number)

            render_passed = (
                source_pages == candidate_pages == 11
                and len(source_images) == len(candidate_images) == 11
                and render_min >= 0.9985
                and not title_bounds_failures
            )

    criteria.append(
        result_item(
            "R009",
            render_passed,
            f"最低标题外像素相似度={render_min:.6f}，"
            f"越界标题页={title_bounds_failures}；{render_evidence}",
        )
    )

    severe = (
        not core_ready
        or not render_passed
        or bool(full_slide_new)
        or bool(markers)
        or not native_passed
    )
    criteria.append(
        result_item("R011", not severe, f"严重错误={severe}")
    )
    order = {identifier: index for index, identifier in enumerate(SPECS)}
    criteria.sort(key=lambda item: order[item["criterion_id"]])
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    output = args.artifact_dir / OUTPUT_NAME
    if output.is_file():
        evaluation = evaluate(output, args.artifact_dir)
    else:
        evaluation = {
            "task_id": TASK_ID,
            "rubric_version": RUBRIC_VERSION,
            "judge_version": JUDGE_VERSION,
            "artifact_sha256": [],
            "criterion_results": [],
            "reward": 0.0,
            "passed": False,
            "hurdle_failures": ["missing_output"],
        }
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(
        json.dumps(evaluation, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    args.reward_file.write_text(
        f"{evaluation['reward']:.6f}\n", encoding="utf-8"
    )
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
