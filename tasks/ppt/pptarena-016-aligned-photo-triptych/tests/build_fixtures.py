#!/usr/bin/env python3
"""只为 PPTArena Case 16 构造可重放的独立合法解、边界例和负例。"""

from __future__ import annotations

import argparse
import shutil
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Inches


OUTPUT_NAME = "OrderPhotosGroundTruth.pptx"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
EMU_PER_POINT = 12_700
TARGET_NAMES = ("Picture 26", "Picture 4", "Picture 18")


def qn(namespace: str, local: str) -> str:
    return f"{{{namespace}}}{local}"


def output_path(root: Path, case: str) -> Path:
    target = root / case / OUTPUT_NAME
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def serialize(root) -> bytes:
    return etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)


def rewrite_archive(source: Path, output: Path, transform) -> None:
    with zipfile.ZipFile(source) as archive_in, zipfile.ZipFile(output, "w") as archive_out:
        for info in archive_in.infolist():
            archive_out.writestr(info, transform(info.filename, archive_in.read(info.filename)))


def add_background(common_slide, *, mode: str) -> None:
    for existing in common_slide.xpath("./p:bg", namespaces={"p": P_NS}):
        common_slide.remove(existing)
    if mode == "inherited":
        return
    background = etree.Element(qn(P_NS, "bg"))
    properties = etree.SubElement(background, qn(P_NS, "bgPr"))
    if mode == "bg1":
        fill = etree.SubElement(properties, qn(A_NS, "solidFill"))
        etree.SubElement(fill, qn(A_NS, "schemeClr"), val="bg1")
    elif mode == "white":
        fill = etree.SubElement(properties, qn(A_NS, "solidFill"))
        etree.SubElement(fill, qn(A_NS, "srgbClr"), val="FFFFFF")
    elif mode == "black":
        fill = etree.SubElement(properties, qn(A_NS, "solidFill"))
        etree.SubElement(fill, qn(A_NS, "srgbClr"), val="000000")
    elif mode == "gradient":
        fill = etree.SubElement(properties, qn(A_NS, "gradFill"), rotWithShape="1")
        stops = etree.SubElement(fill, qn(A_NS, "gsLst"))
        first = etree.SubElement(stops, qn(A_NS, "gs"), pos="0")
        etree.SubElement(first, qn(A_NS, "srgbClr"), val="FFFFFF")
        second = etree.SubElement(stops, qn(A_NS, "gs"), pos="100000")
        etree.SubElement(second, qn(A_NS, "srgbClr"), val="EEEEEE")
        etree.SubElement(fill, qn(A_NS, "lin"), ang="5400000", scaled="1")
    else:
        raise ValueError(mode)
    etree.SubElement(properties, qn(A_NS, "effectLst"))
    common_slide.insert(0, background)


def picture_by_name(root, name: str):
    nodes = root.xpath(
        f"./p:cSld/p:spTree/p:pic[p:nvPicPr/p:cNvPr/@name='{name}']",
        namespaces={"p": P_NS},
    )
    if len(nodes) != 1:
        raise ValueError(f"未唯一定位图片 {name}")
    return nodes[0]


def set_geometry(picture, *, left: float, top: float, size: float, rotation: float = 0.0) -> None:
    namespaces = {"p": P_NS, "a": A_NS}
    transform = picture.find("./p:spPr/a:xfrm", namespaces)
    offset = transform.find("a:off", namespaces)
    extent = transform.find("a:ext", namespaces)
    offset.set("x", str(round(left * EMU_PER_POINT)))
    offset.set("y", str(round(top * EMU_PER_POINT)))
    extent.set("cx", str(round(size * EMU_PER_POINT)))
    extent.set("cy", str(round(size * EMU_PER_POINT)))
    if abs(rotation) < 1e-9:
        transform.attrib.pop("rot", None)
    else:
        transform.set("rot", str(round(rotation * 60_000)))


def layout_variant(
    source: Path,
    output: Path,
    *,
    positions: dict[str, tuple[float, float, float, float]],
    background: str,
) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name != "ppt/slides/slide2.xml":
            return data
        root = etree.fromstring(data)
        common_slide = root.find(qn(P_NS, "cSld"))
        add_background(common_slide, mode=background)
        for picture_name, values in positions.items():
            left, top, size, rotation = values
            set_geometry(picture_by_name(root, picture_name), left=left, top=top, size=size, rotation=rotation)
        return serialize(root)

    rewrite_archive(source, output, transform)


def golden_mutation(golden: Path, output: Path, mutate) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name not in {"ppt/slides/slide1.xml", "ppt/slides/slide2.xml"}:
            return data
        root = etree.fromstring(data)
        changed = mutate(name, root)
        return serialize(root) if changed else data

    rewrite_archive(golden, output, transform)


def delete_picture(golden: Path, output: Path) -> None:
    def mutate(name: str, root) -> bool:
        if name != "ppt/slides/slide2.xml":
            return False
        picture = picture_by_name(root, "Picture 18")
        picture.getparent().remove(picture)
        return True

    golden_mutation(golden, output, mutate)


def swap_positions(golden: Path, output: Path) -> None:
    def mutate(name: str, root) -> bool:
        if name != "ppt/slides/slide2.xml":
            return False
        first = picture_by_name(root, "Picture 26")
        third = picture_by_name(root, "Picture 18")
        namespaces = {"p": P_NS, "a": A_NS}
        first_offset = first.find("./p:spPr/a:xfrm/a:off", namespaces)
        third_offset = third.find("./p:spPr/a:xfrm/a:off", namespaces)
        first_x, third_x = first_offset.get("x"), third_offset.get("x")
        first_offset.set("x", third_x)
        third_offset.set("x", first_x)
        return True

    golden_mutation(golden, output, mutate)


def alter_geometry(golden: Path, output: Path, picture_name: str, *, top=None, size=None, rotation=None) -> None:
    def mutate(name: str, root) -> bool:
        if name != "ppt/slides/slide2.xml":
            return False
        picture = picture_by_name(root, picture_name)
        namespaces = {"p": P_NS, "a": A_NS}
        xfrm = picture.find("./p:spPr/a:xfrm", namespaces)
        offset = xfrm.find("a:off", namespaces)
        extent = xfrm.find("a:ext", namespaces)
        if top is not None:
            offset.set("y", str(round(top * EMU_PER_POINT)))
        if size is not None:
            extent.set("cx", str(round(size * EMU_PER_POINT)))
            extent.set("cy", str(round(size * EMU_PER_POINT)))
        if rotation is not None:
            xfrm.set("rot", str(round(rotation * 60_000)))
        return True

    golden_mutation(golden, output, mutate)


def change_background(golden: Path, output: Path, mode: str) -> None:
    def mutate(name: str, root) -> bool:
        if name != "ppt/slides/slide2.xml":
            return False
        add_background(root.find(qn(P_NS, "cSld")), mode=mode)
        return True

    golden_mutation(golden, output, mutate)


def change_text(golden: Path, output: Path, *, slide_number: int, old: str, new: str) -> None:
    def mutate(name: str, root) -> bool:
        if name != f"ppt/slides/slide{slide_number}.xml":
            return False
        nodes = root.xpath(".//a:t", namespaces={"a": A_NS})
        for node in nodes:
            if node.text and old in node.text:
                node.text = node.text.replace(old, new)
                return True
        raise ValueError(f"未找到文字 {old}")

    golden_mutation(golden, output, mutate)


def move_panel(golden: Path, output: Path) -> None:
    def mutate(name: str, root) -> bool:
        if name != "ppt/slides/slide2.xml":
            return False
        nodes = root.xpath(
            "./p:cSld/p:spTree/p:sp[p:nvSpPr/p:cNvPr/@id='28']/p:spPr/a:xfrm/a:off",
            namespaces={"p": P_NS, "a": A_NS},
        )
        nodes[0].set("x", str(int(nodes[0].get("x")) + 10 * EMU_PER_POINT))
        return True

    golden_mutation(golden, output, mutate)


def recolor_decoration(golden: Path, output: Path) -> None:
    def mutate(name: str, root) -> bool:
        if name != "ppt/slides/slide2.xml":
            return False
        nodes = root.xpath(
            "./p:cSld/p:spTree/p:sp[p:nvSpPr/p:cNvPr/@id='22']/p:spPr",
            namespaces={"p": P_NS, "a": A_NS},
        )
        properties = nodes[0]
        fill = etree.Element(qn(A_NS, "solidFill"))
        etree.SubElement(fill, qn(A_NS, "srgbClr"), val="FF0000")
        line = properties.find(qn(A_NS, "ln"))
        properties.insert(properties.index(line), fill)
        return True

    golden_mutation(golden, output, mutate)


def change_crop(golden: Path, output: Path) -> None:
    def mutate(name: str, root) -> bool:
        if name != "ppt/slides/slide2.xml":
            return False
        picture = picture_by_name(root, "Picture 4")
        source_rect = picture.find("./p:blipFill/a:srcRect", {"p": P_NS, "a": A_NS})
        source_rect.set("l", str(int(source_rect.get("l", "0")) + 5000))
        return True

    golden_mutation(golden, output, mutate)


def swap_picture_content(golden: Path, output: Path) -> None:
    def mutate(name: str, root) -> bool:
        if name != "ppt/slides/slide2.xml":
            return False
        first = picture_by_name(root, "Picture 26").find("./p:blipFill/a:blip", {"p": P_NS, "a": A_NS})
        third = picture_by_name(root, "Picture 18").find("./p:blipFill/a:blip", {"p": P_NS, "a": A_NS})
        attribute = qn(R_NS, "embed")
        first_id, third_id = first.get(attribute), third.get(attribute)
        first.set(attribute, third_id)
        third.set(attribute, first_id)
        return True

    golden_mutation(golden, output, mutate)


def repacked_content_types(golden: Path, output: Path) -> None:
    rewrite_archive(
        golden,
        output,
        lambda name, data: data.replace(b"</Types>", b"  \n</Types>") if name == "[Content_Types].xml" else data,
    )


def add_textbox(golden: Path, output: Path, text: str) -> None:
    presentation = Presentation(golden)
    box = presentation.slides[1].shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(2.5), Inches(0.3))
    box.text = text
    presentation.save(output)


def extra_slide(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    presentation.slides.add_slide(presentation.slide_layouts[0])
    presentation.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--upstream-ground-truth", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()

    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))
    shutil.copy2(args.upstream_ground_truth, output_path(args.outdir, "upstream-ground-truth-unfit"))

    layout_variant(
        args.source,
        output_path(args.outdir, "valid-alt-layout"),
        positions={
            "Picture 26": (76.0, 134.5, 225.5, 0.0),
            "Picture 4": (355.0, 134.5, 225.5, 0.0),
            "Picture 18": (632.0, 134.5, 225.5, 0.0),
        },
        background="white",
    )
    layout_variant(
        args.source,
        output_path(args.outdir, "valid-226-square"),
        positions={
            "Picture 26": (75.0, 134.0, 226.0, 0.0),
            "Picture 4": (354.39, 134.0, 226.0, 0.0),
            "Picture 18": (632.83, 134.0, 226.0, 0.0),
        },
        background="bg1",
    )
    repacked_content_types(args.golden, output_path(args.outdir, "repacked-content-types"))
    delete_picture(args.golden, output_path(args.outdir, "missing-photo"))
    swap_positions(args.golden, output_path(args.outdir, "wrong-order"))
    alter_geometry(args.golden, output_path(args.outdir, "misaligned-top"), "Picture 18", top=145.0)
    alter_geometry(args.golden, output_path(args.outdir, "unequal-size"), "Picture 18", size=210.0)
    alter_geometry(args.golden, output_path(args.outdir, "rotated-waterfall"), "Picture 26", rotation=5.0)
    alter_geometry(args.golden, output_path(args.outdir, "overlap-caption"), "Picture 26", size=230.0)
    change_background(args.golden, output_path(args.outdir, "inherited-background"), "inherited")
    change_background(args.golden, output_path(args.outdir, "gradient-background"), "gradient")
    change_background(args.golden, output_path(args.outdir, "black-background"), "black")
    change_text(args.golden, output_path(args.outdir, "changed-cover-title"), slide_number=1, old="Genetics", new="Genetics X")
    change_text(args.golden, output_path(args.outdir, "changed-label"), slide_number=2, old="WATERFALL", new="WATERFALLS")
    move_panel(args.golden, output_path(args.outdir, "moved-panel"))
    recolor_decoration(args.golden, output_path(args.outdir, "recolored-decoration"))
    change_crop(args.golden, output_path(args.outdir, "changed-crop"))
    swap_picture_content(args.golden, output_path(args.outdir, "swapped-picture-content"))
    change_text(args.golden, output_path(args.outdir, "changed-slide2-title"), slide_number=2, old="Motion & Stills", new="Motion and Stills")
    add_textbox(args.golden, output_path(args.outdir, "added-object"), "extra object")
    extra_slide(args.golden, output_path(args.outdir, "extra-slide"))
    add_textbox(args.golden, output_path(args.outdir, "hardcoded-marker"), "solution/Golden")


if __name__ == "__main__":
    main()
