#!/usr/bin/env python3
"""只为 PPTArena Case 21 构造可重放的有效变体、负例和中等例。"""

from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches


OUTPUT_NAME = "FixTitleCase_GroundTruthA.pptx"
TITLE_IDS = [54, 62, 70, 76, 82, 89, 95, 102, 109, 115, 121, 127, 133, 139, 145]


def output_path(root: Path, case: str) -> Path:
    target = root / case / OUTPUT_NAME
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def audited_title(presentation: Presentation, slide_number: int):
    slide = presentation.slides[slide_number - 1]
    shape_id = TITLE_IDS[slide_number - 1]
    matches = [shape for shape in slide.shapes if shape.shape_id == shape_id]
    if len(matches) != 1 or not getattr(matches[0], "has_text_frame", False):
        raise ValueError(f"第 {slide_number} 页人工审计标题对象缺失")
    return matches[0]


def replace_title_preserving_runs(
    presentation: Presentation,
    slide_number: int,
    replacement: str,
) -> None:
    title = audited_title(presentation, slide_number)
    runs = [run for paragraph in title.text_frame.paragraphs for run in paragraph.runs]
    original = "".join(run.text for run in runs)
    if len(original) != len(replacement):
        raise ValueError(
            f"第 {slide_number} 页替换文本长度不同：{len(original)} != {len(replacement)}"
        )
    offset = 0
    for run in runs:
        width = len(run.text)
        run.text = replacement[offset : offset + width]
        offset += width


def valid_strict_questions(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    replace_title_preserving_runs(presentation, 2, "Who Are We?")
    replace_title_preserving_runs(presentation, 4, "What Do We Do?")
    presentation.save(output)


def missed_one_title(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    replace_title_preserving_runs(presentation, 11, "results of workshops")
    presentation.save(output)


def lowercase_dsi(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    replace_title_preserving_runs(presentation, 12, "Growth of dsi")
    presentation.save(output)


def wrong_question_case(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    replace_title_preserving_runs(presentation, 2, "Who Are we?")
    presentation.save(output)


def wrong_minor_word(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    replace_title_preserving_runs(presentation, 14, "Results Of Growth")
    presentation.save(output)


def wrong_title_style(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    title = audited_title(presentation, 8)
    for paragraph in title.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.text.strip():
                run.font.color.rgb = RGBColor(255, 0, 0)
                presentation.save(output)
                return
    raise ValueError("第 8 页标题没有可变异文字运行")


def changed_body(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    slide = presentation.slides[1]
    title_element = audited_title(presentation, 2)._element
    for shape in slide.shapes:
        if shape._element is title_element or not getattr(shape, "has_text_frame", False):
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.text.strip():
                    run.text += " X"
                    presentation.save(output)
                    return
    raise ValueError("第 2 页没有可变异正文")


def deleted_title(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    title = audited_title(presentation, 10)
    title._element.getparent().remove(title._element)
    presentation.save(output)


def flattened_slide(golden: Path, slide_image: Path, output: Path) -> None:
    presentation = Presentation(golden)
    slide = presentation.slides[9]
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    slide.shapes.add_picture(
        str(slide_image),
        0,
        0,
        width=presentation.slide_width,
        height=presentation.slide_height,
    )
    presentation.save(output)


def hardcoded_marker(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    box = presentation.slides[0].shapes.add_textbox(
        Inches(0.1), Inches(7.1), Inches(2.5), Inches(0.2)
    )
    box.text = "solution/Golden"
    presentation.save(output)


def extra_slide(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(output)


def moved_title(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    title = audited_title(presentation, 10)
    title.top += Inches(0.08)
    presentation.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--slide10-image", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()

    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))
    valid_strict_questions(
        args.golden, output_path(args.outdir, "valid-strict-questions")
    )
    missed_one_title(args.golden, output_path(args.outdir, "missed-one-title"))
    lowercase_dsi(args.golden, output_path(args.outdir, "lowercase-dsi"))
    wrong_question_case(
        args.golden, output_path(args.outdir, "wrong-question-case")
    )
    wrong_minor_word(args.golden, output_path(args.outdir, "wrong-minor-word"))
    wrong_title_style(args.golden, output_path(args.outdir, "wrong-title-style"))
    changed_body(args.golden, output_path(args.outdir, "changed-body"))
    deleted_title(args.golden, output_path(args.outdir, "deleted-title"))
    flattened_slide(
        args.golden,
        args.slide10_image,
        output_path(args.outdir, "flattened"),
    )
    hardcoded_marker(args.golden, output_path(args.outdir, "hardcoded-marker"))
    extra_slide(args.golden, output_path(args.outdir, "extra-slide"))
    moved_title(args.golden, output_path(args.outdir, "moved-title"))


if __name__ == "__main__":
    main()
