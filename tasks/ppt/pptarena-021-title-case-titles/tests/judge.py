#!/usr/bin/env python3
"""PPTArena Case 21 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import math
import re
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from lxml import etree
from PIL import Image, ImageChops, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "FixTitleCase_GroundTruthA.pptx"
INPUT_NAME = "FixTitleCase_TestA.pptx"
INPUT_SHA256 = "5462ea413b6ecdf5ee13b0893c360e744886283949fed7a8b44fd78ed01e727f"
TASK_ID = "cowork-evo/pptarena-021-title-case-titles"
RUBRIC_VERSION = "pptarena-case-021-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-021-deterministic-v0.1"
EMU_PER_INCH = 914_400
GEOMETRY_TOLERANCE = int(0.012 * EMU_PER_INCH)
PAGE_EDGE_TOLERANCE = int(0.03 * EMU_PER_INCH)
SUSPICIOUS_MARKERS = [
    "solution/golden",
    "/root/",
    "judge-result",
    "artifact_dir",
    "pptarena-case-021",
]

# 这些目标来自逐页视觉检查和输入/人工 Ground Truth 的对象级差异审计。
# 不能只调用 slide.shapes.title：前 14 页中的目标标题都是普通文本框。
TITLE_SPECS = [
    {
        "source": "data science and informatics: creating student led data science curricula",
        "allowed": ["Data Science and Informatics: Creating Student Led Data Science Curricula"],
        "geometry": (312600, 2643825, 8518800, 1869300),
        "shape_type": 17,
        "placeholder": "",
    },
    {
        "source": "who are we?",
        "allowed": ["Who are we?", "Who Are We?"],
        "geometry": (533400, 457200, 6477000, 769441),
        "shape_type": 17,
        "placeholder": "",
    },
    {
        "source": "brief history of dsi",
        "allowed": ["Brief History of DSI"],
        "geometry": (533400, 457200, 6477000, 769500),
        "shape_type": 17,
        "placeholder": "",
    },
    {
        "source": "what do we do?",
        "allowed": ["What do we do?", "What Do We Do?"],
        "geometry": (533400, 457200, 6477000, 769500),
        "shape_type": 17,
        "placeholder": "",
    },
    {
        "source": "our workshops:",
        "allowed": ["Our Workshops:"],
        "geometry": (533400, 457200, 6477000, 769500),
        "shape_type": 17,
        "placeholder": "",
    },
    {
        "source": "workshop Coordination",
        "allowed": ["Workshop Coordination"],
        "geometry": (533400, 457200, 6477000, 769500),
        "shape_type": 17,
        "placeholder": "",
    },
    {
        "source": "competing data science truths ",
        "allowed": ["Competing Data Science Truths"],
        "geometry": (533400, 457200, 6477000, 769500),
        "shape_type": 17,
        "placeholder": "",
    },
    {
        "source": "Workshop design\t",
        "allowed": ["Workshop Design"],
        "geometry": (533400, 457200, 6477000, 769500),
        "shape_type": 17,
        "placeholder": "",
    },
    {
        "source": "workshop instructors",
        "allowed": ["Workshop Instructors"],
        "geometry": (533400, 457200, 6477000, 769500),
        "shape_type": 17,
        "placeholder": "",
    },
    {
        "source": "workshop instructors",
        "allowed": ["Workshop Instructors"],
        "geometry": (533400, 457200, 6477000, 769500),
        "shape_type": 17,
        "placeholder": "",
    },
    {
        "source": "results of workshops",
        "allowed": ["Results of Workshops"],
        "geometry": (533400, 457200, 6477000, 769500),
        "shape_type": 17,
        "placeholder": "",
    },
    {
        "source": "growth of dsi",
        "allowed": ["Growth of DSI"],
        "geometry": (533400, 457200, 6477000, 769500),
        "shape_type": 17,
        "placeholder": "",
    },
    {
        "source": "Growth Statistics",
        "allowed": ["Growth Statistics"],
        "geometry": (533400, 457200, 6477000, 769500),
        "shape_type": 17,
        "placeholder": "",
    },
    {
        "source": "results of growth",
        "allowed": ["Results of Growth"],
        "geometry": (533400, 457200, 6477000, 769500),
        "shape_type": 17,
        "placeholder": "",
    },
    {
        "source": "contact",
        "allowed": ["Contact"],
        "geometry": (685800, 838150, 7772400, 1470000),
        "shape_type": 14,
        "placeholder": "CENTER_TITLE (3)",
    },
]

SPECS = {
    "R001": ("指定 PPTX 可打开，严格保留 15 页与页面尺寸", 6, "hurdle"),
    "R002": ("第 1–15 页语义主标题均采用允许的 Title Case", 16, "hurdle"),
    "R003": ("DSI 与两道问句标题采用允许写法", 8, "hurdle"),
    "R004": ("标题除字母大小写外的内容和样式保持", 9, "hurdle"),
    "R005": ("全部标题的文本框语义和几何保持", 7, "positive"),
    "R006": ("15 页非标题原文与讲者备注按页完整保持", 9, "hurdle"),
    "R007": ("全部非标题文字对象的样式、几何和层级保持", 8, "hurdle"),
    "R008": ("图片和非文本图形保持且未压平", 6, "hurdle"),
    "R009": ("15 页可渲染且标题外视觉保持", 9, "positive"),
    "R010": ("输出目录干净且不含答案路径或未授权文字", 4, "positive"),
    "R011": ("页数、标题、缩写、非标题内容或可编辑性无严重错误", -20, "penalty"),
}


def normalize_visible(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def locate_source() -> Path | None:
    candidates = [
        Path("/app/input_files") / INPUT_NAME,
        Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME,
    ]
    return next((path for path in candidates if path.is_file()), None)


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def txbody(shape):
    if shape is None:
        return None
    nodes = shape._element.xpath('./*[local-name()="txBody"]')
    return nodes[0] if nodes else None


def shape_has_txbody(shape) -> bool:
    return txbody(shape) is not None


def shape_text(shape) -> str:
    body = txbody(shape)
    if body is None:
        return ""
    return "\n".join(
        "".join(paragraph.xpath('.//*[local-name()="t"]/text()'))
        for paragraph in body.xpath('./*[local-name()="p"]')
    )


def geometry(shape) -> tuple[int, int, int, int]:
    return (int(shape.left), int(shape.top), int(shape.width), int(shape.height))


def rounded_geometry(shape) -> tuple[int, int, int, int]:
    return tuple(round(value / 1000) * 1000 for value in geometry(shape))


def geometry_distance(shape, expected: tuple[int, int, int, int]) -> int:
    return sum(abs(actual - target) for actual, target in zip(geometry(shape), expected))


def placeholder_type(shape) -> str:
    if shape is None or not getattr(shape, "is_placeholder", False):
        return ""
    return str(shape.placeholder_format.type)


def title_shape(slide, spec: dict):
    candidates = [shape for shape in iter_shapes(slide.shapes) if shape_has_txbody(shape)]
    if not candidates:
        return None
    source_key = normalize_visible(spec["source"]).casefold()
    same_letters = [
        shape
        for shape in candidates
        if normalize_visible(shape_text(shape)).casefold() == source_key
    ]
    pool = same_letters or candidates
    return min(pool, key=lambda shape: geometry_distance(shape, spec["geometry"]))


def local_attr_name(name: str) -> str:
    return etree.QName(name).localname if name.startswith("{") else name


def canonical_txbody(shape) -> bytes:
    body = txbody(shape)
    if body is None:
        return b""
    clone = copy.deepcopy(body)
    for node in clone.iter():
        for attribute in list(node.attrib):
            if local_attr_name(attribute) in {"dirty", "smtClean"}:
                del node.attrib[attribute]
    return etree.tostring(clone, method="c14n")


def canonical_title_style_shell(shape) -> bytes:
    body = txbody(shape)
    if body is None:
        return b""
    clone = copy.deepcopy(body)
    for node in list(
        clone.xpath('.//*[local-name()="r" or local-name()="fld" or local-name()="br"]')
    ):
        node.getparent().remove(node)
    for node in clone.iter():
        for attribute in list(node.attrib):
            if local_attr_name(attribute) in {"dirty", "smtClean"}:
                del node.attrib[attribute]
    return etree.tostring(clone, method="c14n")


def color_signature(font) -> tuple:
    color = font.color
    try:
        color_type = int(color.type) if color.type is not None else None
    except Exception:
        color_type = str(color.type)
    try:
        rgb = str(color.rgb) if color.rgb is not None else None
    except Exception:
        rgb = None
    try:
        theme = int(color.theme_color) if color.theme_color is not None else None
    except Exception:
        theme = None
    try:
        brightness = round(color.brightness, 5)
    except Exception:
        brightness = None
    return (color_type, rgb, theme, brightness)


def run_style_set(shape) -> set[tuple]:
    if shape is None or not getattr(shape, "has_text_frame", False):
        return set()
    styles: set[tuple] = set()
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            font = run.font
            styles.add(
                (
                    font.name,
                    round(font.size.pt, 4) if font.size else None,
                    bool(font.bold),
                    bool(font.italic),
                    bool(font.underline),
                    int(font.language_id) if font.language_id is not None else None,
                    color_signature(font),
                )
            )
    return styles


def text_shape_record(shape) -> tuple:
    return (
        int(shape.shape_type),
        placeholder_type(shape),
        normalize_visible(shape_text(shape)),
        rounded_geometry(shape),
        hashlib.sha256(canonical_txbody(shape)).hexdigest(),
    )


def non_title_text_records(slide, title) -> list[tuple]:
    title_element = title._element if title is not None else None
    return [
        text_shape_record(shape)
        for shape in iter_shapes(slide.shapes)
        if shape._element is not title_element and shape_has_txbody(shape)
    ]


def non_title_text_counter(slide, title) -> Counter[str]:
    title_element = title._element if title is not None else None
    values: Counter[str] = Counter()
    for shape in iter_shapes(slide.shapes):
        if shape._element is title_element or not shape_has_txbody(shape):
            continue
        text = normalize_visible(shape_text(shape))
        if text:
            values[text] += 1
    return values


def slide_notes(slide) -> tuple[str, ...]:
    try:
        return tuple(
            shape.text
            for shape in slide.notes_slide.shapes
            if getattr(shape, "has_text_frame", False)
        )
    except Exception:
        return ()


def slide_picture_counter(slide) -> Counter[tuple]:
    values: Counter[tuple] = Counter()
    for shape in iter_shapes(slide.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            values[(hashlib.sha256(shape.image.blob).hexdigest(), rounded_geometry(shape))] += 1
    return values


def non_text_shape_records(slide) -> list[tuple]:
    values: list[tuple] = []
    for shape in iter_shapes(slide.shapes):
        if shape_has_txbody(shape):
            continue
        picture_hash = ""
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            picture_hash = hashlib.sha256(shape.image.blob).hexdigest()
        values.append((int(shape.shape_type), rounded_geometry(shape), picture_hash))
    return values


def media_hashes(path: Path) -> Counter[str]:
    with zipfile.ZipFile(path) as archive:
        return Counter(
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("ppt/media/")
        )


def package_ok(path: Path) -> tuple[bool, str]:
    try:
        with zipfile.ZipFile(path) as archive:
            broken = archive.testzip()
            names = set(archive.namelist())
            required = {
                "[Content_Types].xml",
                "ppt/presentation.xml",
                "ppt/slides/slide1.xml",
                "ppt/slides/slide15.xml",
            }
            missing = sorted(required - names)
            macros = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
        return broken is None and not missing and not macros, (
            f"损坏成员={broken}，缺失={missing}，宏={macros}"
        )
    except Exception as exc:
        return False, f"ZIP 检查失败：{exc}"


def render_presentation(path: Path, outdir: Path) -> tuple[list[Path], int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "lo-profile"
    profile.mkdir()
    completed = subprocess.run(
        [
            "soffice",
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(outdir),
            str(path),
        ],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return [], 0, f"转 PDF 失败：{completed.stdout[-500:]}"
    info = subprocess.run(
        ["pdfinfo", str(pdf)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=30,
        check=False,
    ).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    pages = int(match.group(1)) if match else 0
    prefix = outdir / "slide"
    raster = subprocess.run(
        ["pdftoppm", "-jpeg", "-r", "90", str(pdf), str(prefix)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    images = sorted(outdir.glob("slide-*.jpg")) if raster.returncode == 0 else []
    return images, pages, (
        f"页数={pages}，渲染图={len(images)}，"
        f"soffice={completed.returncode}，pdftoppm={raster.returncode}"
    )


def emu_rect_to_pixels(
    shape,
    image_size: tuple[int, int],
    slide_size: tuple[int, int],
    # 第 1 页标题本身已经溢出文本框并与作者行重叠；128 像素只用于
    # 排除标题大小写导致的字形差异，正文仍由结构判据逐对象保护。
    padding: int = 128,
) -> tuple[int, int, int, int]:
    image_width, image_height = image_size
    slide_width, slide_height = slide_size
    left, top, width, height = geometry(shape)
    return (
        max(0, math.floor(left * image_width / slide_width) - padding),
        max(0, math.floor(top * image_height / slide_height) - padding),
        min(image_width, math.ceil((left + width) * image_width / slide_width) + padding),
        min(image_height, math.ceil((top + height) * image_height / slide_height) + padding),
    )


def page_overflow(shape, slide_size: tuple[int, int]) -> tuple[int, int, int, int]:
    slide_width, slide_height = slide_size
    left, top, width, height = geometry(shape)
    return (
        max(0, -left),
        max(0, -top),
        max(0, left + width - slide_width),
        max(0, top + height - slide_height),
    )


def image_similarity(
    first: Path,
    second: Path,
    ignored_rectangles: list[tuple[int, int, int, int]],
) -> float:
    with Image.open(first) as left_image, Image.open(second) as right_image:
        left = left_image.convert("RGB")
        right = right_image.convert("RGB")
        if left.size != right.size:
            return 0.0
        difference = ImageChops.difference(left, right)
        mask = Image.new("L", left.size, 255)
        draw = ImageDraw.Draw(mask)
        for rectangle in ignored_rectangles:
            draw.rectangle(rectangle, fill=0)
        rgb_mask = Image.merge("RGB", (mask, mask, mask))
        masked = ImageChops.multiply(difference, rgb_mask)
        histogram = masked.histogram()
        squared = sum(count * ((index % 256) ** 2) for index, count in enumerate(histogram))
        unmasked_pixels = mask.histogram()[255]
        if unmasked_pixels <= 0:
            return 0.0
        rms = math.sqrt(squared / (unmasked_pixels * 3))
        return max(0.0, 1.0 - rms / 255.0)


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {
        "criterion_id": identifier,
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
        "raw": {
            "description": description,
            "weight": weight,
            "type": kind,
            "passed": bool(passed),
        },
    }


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [
        item["criterion_id"]
        for item in criteria
        if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]
    ]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(
        item["raw"]["weight"]
        for item in criteria
        if item["raw"]["weight"] > 0 and item["raw"]["passed"]
    )
    penalty = sum(
        abs(item["raw"]["weight"])
        for item in criteria
        if item["raw"]["weight"] < 0 and not item["raw"]["passed"]
    )
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [sha256(path)] if path.is_file() else [],
        "criterion_results": criteria,
        "reward": reward,
        "passed": reward >= 0.8 and not hurdles,
        "hurdle_failures": hurdles,
    }


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        criteria = [
            result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配")
            for identifier in SPECS
        ]
        return make_result(path, criteria)
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        criteria = [
            result_item(identifier, False, f"PPTX 打开失败：{exc}")
            for identifier in SPECS
        ]
        return make_result(path, criteria)

    criteria: list[dict] = []
    package_passed, package_evidence = package_ok(path)
    size_passed = (
        len(candidate.slides) == len(source.slides) == len(TITLE_SPECS)
        and candidate.slide_width == source.slide_width
        and candidate.slide_height == source.slide_height
    )
    criteria.append(
        result_item(
            "R001",
            package_passed and size_passed,
            f"{package_evidence}，页数={len(candidate.slides)}，"
            f"页面={candidate.slide_width / EMU_PER_INCH:.4f}×"
            f"{candidate.slide_height / EMU_PER_INCH:.4f}英寸",
        )
    )

    source_titles = []
    candidate_titles = []
    title_failures: list[str] = []
    semantic_failures: list[str] = []
    style_failures: list[int] = []
    geometry_failures: list[int] = []

    if len(candidate.slides) == len(source.slides) == len(TITLE_SPECS):
        for slide_number, (source_slide, candidate_slide, spec) in enumerate(
            zip(source.slides, candidate.slides, TITLE_SPECS), start=1
        ):
            source_title = title_shape(source_slide, spec)
            candidate_title = title_shape(candidate_slide, spec)
            source_titles.append(source_title)
            candidate_titles.append(candidate_title)
            if source_title is None or candidate_title is None:
                title_failures.append(f"第{slide_number}页缺失语义主标题")
                style_failures.append(slide_number)
                geometry_failures.append(slide_number)
                continue

            actual = normalize_visible(shape_text(candidate_title))
            if actual not in spec["allowed"]:
                title_failures.append(
                    f"第{slide_number}页={actual!r}，允许={spec['allowed']}"
                )
            if slide_number in {3, 12} and "DSI" not in actual:
                semantic_failures.append(f"第{slide_number}页没有保留 DSI")
            if slide_number in {2, 4} and actual not in spec["allowed"]:
                semantic_failures.append(f"第{slide_number}页问句写法不在允许集合")

            source_letters = normalize_visible(shape_text(source_title)).casefold()
            candidate_letters = actual.casefold()
            style_ok = (
                candidate_letters == source_letters
                and canonical_title_style_shell(candidate_title)
                == canonical_title_style_shell(source_title)
                and run_style_set(candidate_title) == run_style_set(source_title)
            )
            if not style_ok:
                style_failures.append(slide_number)

            geometry_ok = all(
                abs(actual_value - expected_value) <= GEOMETRY_TOLERANCE
                for actual_value, expected_value in zip(
                    geometry(candidate_title), geometry(source_title)
                )
            )
            semantic_type_ok = (
                int(candidate_title.shape_type) == spec["shape_type"]
                and placeholder_type(candidate_title) == spec["placeholder"]
            )
            if not geometry_ok or not semantic_type_ok:
                geometry_failures.append(slide_number)
    else:
        title_failures.append("页数不等于 15，无法逐页检查标题")
        semantic_failures.append("页数不等于 15，无法检查 DSI 和问句")
        style_failures = list(range(1, 16))
        geometry_failures = list(range(1, 16))

    criteria.append(result_item("R002", not title_failures, f"标题失败={title_failures}"))
    criteria.append(
        result_item("R003", not semantic_failures, f"语义失败={semantic_failures}")
    )
    criteria.append(
        result_item(
            "R004",
            not style_failures,
            f"标题非大小写文字或样式变化页={sorted(set(style_failures))}",
        )
    )
    criteria.append(
        result_item(
            "R005",
            not geometry_failures,
            f"标题几何、文本框或占位符变化页={sorted(set(geometry_failures))}",
        )
    )

    source_non_title_texts = [
        non_title_text_counter(slide, title)
        for slide, title in zip(source.slides, source_titles)
    ]
    candidate_non_title_texts = [
        non_title_text_counter(slide, title)
        for slide, title in zip(candidate.slides, candidate_titles)
    ]
    text_mismatches = (
        [
            index + 1
            for index, (actual, expected) in enumerate(
                zip(candidate_non_title_texts, source_non_title_texts)
            )
            if actual != expected
        ]
        if len(candidate_non_title_texts) == len(source_non_title_texts) == 15
        else list(range(1, 16))
    )
    source_notes = [slide_notes(slide) for slide in source.slides]
    candidate_notes = [slide_notes(slide) for slide in candidate.slides]
    notes_mismatches = (
        [
            index + 1
            for index, (actual, expected) in enumerate(zip(candidate_notes, source_notes))
            if actual != expected
        ]
        if len(candidate_notes) == len(source_notes) == 15
        else list(range(1, 16))
    )
    criteria.append(
        result_item(
            "R006",
            not text_mismatches and not notes_mismatches,
            f"非标题文字不一致页={text_mismatches}，备注不一致页={notes_mismatches}",
        )
    )

    non_title_mismatches = (
        [
            index + 1
            for index, (candidate_slide, source_slide, candidate_title, source_title) in enumerate(
                zip(candidate.slides, source.slides, candidate_titles, source_titles)
            )
            if non_title_text_records(candidate_slide, candidate_title)
            != non_title_text_records(source_slide, source_title)
        ]
        if len(candidate.slides) == len(source.slides) == 15
        else list(range(1, 16))
    )
    criteria.append(
        result_item(
            "R007",
            not non_title_mismatches,
            f"非标题文字对象变化页={non_title_mismatches}",
        )
    )

    picture_mismatches = (
        [
            index + 1
            for index, (candidate_slide, source_slide) in enumerate(
                zip(candidate.slides, source.slides)
            )
            if slide_picture_counter(candidate_slide) != slide_picture_counter(source_slide)
        ]
        if len(candidate.slides) == len(source.slides) == 15
        else list(range(1, 16))
    )
    non_text_mismatches = (
        [
            index + 1
            for index, (candidate_slide, source_slide) in enumerate(
                zip(candidate.slides, source.slides)
            )
            if non_text_shape_records(candidate_slide) != non_text_shape_records(source_slide)
        ]
        if len(candidate.slides) == len(source.slides) == 15
        else list(range(1, 16))
    )
    full_slide_new: list[int] = []
    for slide_index, slide in enumerate(candidate.slides, start=1):
        source_pictures = (
            slide_picture_counter(source.slides[slide_index - 1])
            if slide_index <= len(source.slides)
            else Counter()
        )
        for shape in iter_shapes(slide.shapes):
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            coverage = shape.width * shape.height / (
                candidate.slide_width * candidate.slide_height
            )
            key = (hashlib.sha256(shape.image.blob).hexdigest(), rounded_geometry(shape))
            if coverage >= 0.9 and source_pictures[key] == 0:
                full_slide_new.append(slide_index)
    try:
        media_equal = media_hashes(path) == media_hashes(source_path)
    except Exception:
        media_equal = False
    native_passed = (
        not picture_mismatches
        and not non_text_mismatches
        and not full_slide_new
        and media_equal
    )
    criteria.append(
        result_item(
            "R008",
            native_passed,
            f"图片变化页={picture_mismatches}，非文本对象变化页={non_text_mismatches}，"
            f"新增整页图片页={full_slide_new}，媒体哈希集合一致={media_equal}",
        )
    )

    artifact_files = sorted(
        str(item.relative_to(artifact_dir))
        for item in artifact_dir.rglob("*")
        if item.is_file()
    )
    all_text = "\n".join(
        shape_text(shape)
        for slide in candidate.slides
        for shape in iter_shapes(slide.shapes)
    )
    all_notes = "\n".join(text for slide in candidate.slides for text in slide_notes(slide))
    lowered = f"{all_text}\n{all_notes}".lower()
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker.lower() in lowered]
    hygiene_passed = artifact_files == [OUTPUT_NAME] and not markers
    criteria.append(
        result_item(
            "R010",
            hygiene_passed,
            f"输出文件={artifact_files}，可疑标记={markers}",
        )
    )

    core_ready = all(
        next(item for item in criteria if item["criterion_id"] == identifier)["raw"]["passed"]
        for identifier in ["R001", "R002", "R003", "R004", "R006", "R007", "R008"]
    )
    render_passed = False
    render_min = 0.0
    title_bounds_failures: list[int] = []
    render_evidence = "前置结构或内容检查失败，未执行渲染"

    if core_ready:
        with tempfile.TemporaryDirectory(prefix="pptarena021-judge-") as temporary:
            root = Path(temporary)
            source_images, source_pages, source_note = render_presentation(
                source_path, root / "source"
            )
            candidate_images, candidate_pages, candidate_note = render_presentation(
                path, root / "candidate"
            )
            render_evidence = f"源[{source_note}]；候选[{candidate_note}]"
            scores: list[float] = []
            if len(source_images) == len(candidate_images) == 15:
                for index, (source_image, candidate_image) in enumerate(
                    zip(source_images, candidate_images)
                ):
                    with Image.open(source_image) as image:
                        image_size = image.size
                    ignored: list[tuple[int, int, int, int]] = []
                    source_title = source_titles[index]
                    candidate_title = candidate_titles[index]
                    if source_title is not None:
                        ignored.append(
                            emu_rect_to_pixels(
                                source_title,
                                image_size,
                                (source.slide_width, source.slide_height),
                            )
                        )
                    if candidate_title is not None:
                        ignored.append(
                            emu_rect_to_pixels(
                                candidate_title,
                                image_size,
                                (candidate.slide_width, candidate.slide_height),
                            )
                        )
                    scores.append(image_similarity(source_image, candidate_image, ignored))
            render_min = min(scores) if scores else 0.0

            for slide_number, shape in enumerate(candidate_titles, start=1):
                if shape is None:
                    title_bounds_failures.append(slide_number)
                    continue
                source_shape = source_titles[slide_number - 1]
                if source_shape is None:
                    title_bounds_failures.append(slide_number)
                    continue
                source_overflow = page_overflow(
                    source_shape, (source.slide_width, source.slide_height)
                )
                candidate_overflow = page_overflow(
                    shape, (candidate.slide_width, candidate.slide_height)
                )
                if any(
                    actual > expected + PAGE_EDGE_TOLERANCE
                    for actual, expected in zip(candidate_overflow, source_overflow)
                ):
                    title_bounds_failures.append(slide_number)

            render_passed = (
                source_pages == candidate_pages == 15
                and len(source_images) == len(candidate_images) == 15
                and render_min >= 0.9985
                and not title_bounds_failures
            )

    criteria.append(
        result_item(
            "R009",
            render_passed,
            f"最低标题外像素相似度={render_min:.6f}，越界标题页={title_bounds_failures}；"
            f"{render_evidence}",
        )
    )

    severe = (
        not core_ready
        or not render_passed
        or bool(full_slide_new)
        or bool(markers)
        or not native_passed
    )
    criteria.append(result_item("R011", not severe, f"严重错误={severe}"))
    order = {identifier: index for index, identifier in enumerate(SPECS)}
    criteria.sort(key=lambda item: order[item["criterion_id"]])
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    output = args.artifact_dir / OUTPUT_NAME
    if output.is_file():
        evaluation = evaluate(output, args.artifact_dir)
    else:
        evaluation = {
            "task_id": TASK_ID,
            "rubric_version": RUBRIC_VERSION,
            "judge_version": JUDGE_VERSION,
            "artifact_sha256": [],
            "criterion_results": [],
            "reward": 0.0,
            "passed": False,
            "hurdle_failures": ["missing_output"],
        }
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(
        json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
