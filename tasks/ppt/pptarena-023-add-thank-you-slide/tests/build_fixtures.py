#!/usr/bin/env python3
"""只为 PPTArena Case 23 构造可重放的有效变体、负例和中等例。"""

from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUTPUT_NAME = "AddThankYouSlide_GroundTruthA.pptx"


def output_path(root: Path, case: str) -> Path:
    target = root / case / OUTPUT_NAME
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def exact_thank_you_shape(presentation: Presentation):
    matches = [
        shape
        for shape in presentation.slides[-1].shapes
        if getattr(shape, "has_text_frame", False)
        and " ".join(shape.text.split()) == "Thank you"
    ]
    if len(matches) != 1:
        raise ValueError(f"末页应有唯一 Thank you 文字对象，实际为 {len(matches)}")
    return matches[0]


def append_valid_alternate(source: Path, output: Path) -> None:
    """不复制 Golden 的主题布局，用空白版式和显式文本框实现同义答案。"""
    presentation = Presentation(source)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(
        Inches(1.25), Inches(3.05), Inches(7.50), Inches(0.95)
    )
    box.text_frame.clear()
    box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = box.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = "Thank you"
    run.font.size = Pt(44)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 78, 121)
    presentation.save(output)


def wrong_title(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    exact_thank_you_shape(presentation).text = "Thanks"
    presentation.save(output)


def duplicate_title(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    slide = presentation.slides[-1]
    box = slide.shapes.add_textbox(
        Inches(3.25), Inches(5.50), Inches(3.50), Inches(0.60)
    )
    box.text = "Thank you"
    presentation.save(output)


def inserted_at_start(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    slide_ids = presentation.slides._sldIdLst
    last = slide_ids[-1]
    slide_ids.remove(last)
    slide_ids.insert(0, last)
    presentation.save(output)


def changed_original_body(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    for slide_index in range(14):
        slide = presentation.slides[slide_index]
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip():
                        run.text += " X"
                        presentation.save(output)
                        return
    raise ValueError("原有 14 页没有可变异正文")


def changed_original_chart(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    for slide_index in range(14):
        slide = presentation.slides[slide_index]
        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.CHART:
                continue
            for node in shape.chart._element.xpath('.//*[local-name()="v"]'):
                if not node.text:
                    continue
                try:
                    value = float(node.text)
                except ValueError:
                    continue
                node.text = f"{value + 1:g}"
                presentation.save(output)
                return
    raise ValueError("原有 14 页没有可变异图表数值")


def flattened_final(golden: Path, slide_image: Path, output: Path) -> None:
    presentation = Presentation(golden)
    slide = presentation.slides[-1]
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    slide.shapes.add_picture(
        str(slide_image),
        0,
        0,
        width=presentation.slide_width,
        height=presentation.slide_height,
    )
    presentation.save(output)


def hardcoded_marker(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    box = presentation.slides[-1].shapes.add_textbox(
        Inches(0.10), Inches(7.10), Inches(2.60), Inches(0.20)
    )
    box.text = "solution/Golden"
    presentation.save(output)


def invisible_title(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    title = exact_thank_you_shape(presentation)
    for paragraph in title.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.text.strip():
                run.font.color.rgb = RGBColor(255, 255, 255)
    presentation.save(output)


def off_center_title(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    exact_thank_you_shape(presentation).top += Inches(1.75)
    presentation.save(output)


def extra_slide(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--slide15-image", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()

    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))
    append_valid_alternate(args.source, output_path(args.outdir, "valid-alternate"))
    wrong_title(args.golden, output_path(args.outdir, "wrong-title"))
    duplicate_title(args.golden, output_path(args.outdir, "duplicate-title"))
    inserted_at_start(args.golden, output_path(args.outdir, "inserted-at-start"))
    changed_original_body(args.golden, output_path(args.outdir, "changed-original-body"))
    changed_original_chart(args.golden, output_path(args.outdir, "changed-original-chart"))
    flattened_final(
        args.golden,
        args.slide15_image,
        output_path(args.outdir, "flattened-final"),
    )
    hardcoded_marker(args.golden, output_path(args.outdir, "hardcoded-marker"))
    invisible_title(args.golden, output_path(args.outdir, "invisible-title"))
    off_center_title(args.golden, output_path(args.outdir, "off-center-title"))
    extra_slide(args.golden, output_path(args.outdir, "extra-slide"))


if __name__ == "__main__":
    main()
