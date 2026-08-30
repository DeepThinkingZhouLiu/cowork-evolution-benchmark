#!/usr/bin/env python3
"""PPTArena Case 23 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import statistics
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from PIL import Image, ImageChops, ImageStat
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "AddThankYouSlide_GroundTruthA.pptx"
INPUT_NAME = "AddThankYouSlide_TestA.pptx"
INPUT_SHA256 = "cd6c301ffa1fe987814217e02220e63404c9481e20aebc9d2fb3c3394727ebbc"
TASK_ID = "cowork-evo/pptarena-023-add-thank-you-slide"
RUBRIC_VERSION = "pptarena-case-023-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-023-deterministic-v0.1"
EMU_PER_INCH = 914_400
PAGE_EDGE_TOLERANCE = int(0.03 * EMU_PER_INCH)
SUSPICIOUS_MARKERS = [
    "solution/golden",
    "/root/",
    "judge-result",
    "artifact_dir",
    "pptarena-case-023",
]

SPECS = {
    "R001": ("指定 PPTX 可打开，页面尺寸不变且总页数为 15", 6, "hurdle"),
    "R002": ("原有 14 页内容、顺序、图表、图片和结构保持", 15, "hurdle"),
    "R003": ("原有 14 页样式、几何、备注和媒体保持", 10, "hurdle"),
    "R004": ("第 15 页存在唯一可编辑主标题 Thank you", 15, "hurdle"),
    "R005": ("第 15 页标题醒目且合理居中", 8, "positive"),
    "R006": ("第 15 页简洁且没有无关正文或重复标题", 6, "positive"),
    "R007": ("原有 14 页逐页渲染视觉保持", 8, "hurdle"),
    "R008": ("全部页面保持原生可编辑且没有压平或宏", 7, "hurdle"),
    "R009": ("15 页可渲染且末页标题可见、未越界", 9, "positive"),
    "R010": ("输出目录干净且不含答案路径或未授权文字", 4, "positive"),
    "R011": ("页数、原页、末页标题和可编辑性无严重错误", -20, "penalty"),
}


def normalize_visible(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def locate_source() -> Path | None:
    candidates = [
        Path("/app/input_files") / INPUT_NAME,
        Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME,
    ]
    return next((path for path in candidates if path.is_file()), None)


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def placeholder_type(shape) -> str:
    if not getattr(shape, "is_placeholder", False):
        return ""
    return str(shape.placeholder_format.type)


def geometry(shape) -> tuple[int, int, int, int]:
    return (int(shape.left), int(shape.top), int(shape.width), int(shape.height))


def rounded_geometry(shape) -> tuple[int, int, int, int]:
    return tuple(round(value / 1000) * 1000 for value in geometry(shape))


def color_signature(font) -> tuple:
    color = font.color
    try:
        color_type = int(color.type) if color.type is not None else None
    except Exception:
        color_type = str(color.type)
    try:
        rgb = str(color.rgb) if color.rgb is not None else None
    except Exception:
        rgb = None
    try:
        theme = int(color.theme_color) if color.theme_color is not None else None
    except Exception:
        theme = None
    try:
        brightness = round(color.brightness, 5)
    except Exception:
        brightness = None
    return (color_type, rgb, theme, brightness)


def run_style_signature(font) -> tuple:
    return (
        font.name,
        round(font.size.pt, 4) if font.size else None,
        bool(font.bold),
        bool(font.italic),
        bool(font.underline),
        int(font.language_id) if font.language_id is not None else None,
        color_signature(font),
    )


def text_format_signature(shape) -> tuple:
    text_frame = shape.text_frame
    paragraphs = []
    for paragraph in text_frame.paragraphs:
        styles = {
            run_style_signature(run.font)
            for run in paragraph.runs
            if run.text.strip()
        }
        paragraphs.append(
            (
                normalize_visible(paragraph.text),
                paragraph.level,
                str(paragraph.alignment),
                str(paragraph.line_spacing),
                str(paragraph.space_before),
                str(paragraph.space_after),
                tuple(sorted(styles, key=repr)),
            )
        )
    return (
        int(shape.shape_type),
        placeholder_type(shape),
        rounded_geometry(shape),
        (
            text_frame.margin_left,
            text_frame.margin_right,
            text_frame.margin_top,
            text_frame.margin_bottom,
            str(text_frame.vertical_anchor),
            text_frame.word_wrap,
            str(text_frame.auto_size),
        ),
        tuple(paragraphs),
    )


def chart_semantic_signature(shape) -> tuple:
    chart = shape.chart
    values = tuple(
        normalize_visible(value)
        for value in chart._element.xpath('.//*[local-name()="v"]/text()')
    )
    series = []
    for item in chart.series:
        try:
            item_values = tuple(float(value) for value in item.values)
        except Exception:
            item_values = tuple(str(value) for value in item.values)
        series.append((item.name or "", item_values))
    return (str(chart.chart_type), values, tuple(series))


def shape_content_signature(shape) -> tuple:
    text = normalize_visible(shape.text) if getattr(shape, "has_text_frame", False) else ""
    picture_hash = ""
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        picture_hash = hashlib.sha256(shape.image.blob).hexdigest()
    chart = ()
    if shape.shape_type == MSO_SHAPE_TYPE.CHART:
        chart = chart_semantic_signature(shape)
    children = ()
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        children = tuple(shape_content_signature(child) for child in shape.shapes)
    return (
        str(shape.shape_type),
        placeholder_type(shape),
        text,
        picture_hash,
        chart,
        children,
    )


def shape_style_signature(shape) -> tuple:
    text_style = ()
    if getattr(shape, "has_text_frame", False):
        text_style = text_format_signature(shape)
    picture_hash = ""
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        picture_hash = hashlib.sha256(shape.image.blob).hexdigest()
    children = ()
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        children = tuple(shape_style_signature(child) for child in shape.shapes)
    return (
        str(shape.shape_type),
        placeholder_type(shape),
        rounded_geometry(shape),
        text_style,
        picture_hash,
        children,
    )


def slide_content_signature(slide) -> tuple:
    return tuple(shape_content_signature(shape) for shape in slide.shapes)


def slide_style_signature(slide) -> tuple:
    return tuple(shape_style_signature(shape) for shape in slide.shapes)


def slide_notes(slide) -> tuple[str, ...]:
    try:
        return tuple(
            shape.text
            for shape in slide.notes_slide.shapes
            if getattr(shape, "has_text_frame", False)
        )
    except Exception:
        return ()


def media_hashes(path: Path) -> Counter[str]:
    with zipfile.ZipFile(path) as archive:
        return Counter(
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("ppt/media/")
        )


def package_ok(path: Path) -> tuple[bool, str]:
    try:
        with zipfile.ZipFile(path) as archive:
            broken = archive.testzip()
            names = set(archive.namelist())
            required = {
                "[Content_Types].xml",
                "ppt/presentation.xml",
                "ppt/slides/slide1.xml",
                "ppt/slides/slide15.xml",
            }
            missing = sorted(required - names)
            macros = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
        return broken is None and not missing and not macros, (
            f"损坏成员={broken}，缺失={missing}，宏={macros}"
        )
    except Exception as exc:
        return False, f"ZIP 检查失败：{exc}"


def thank_you_shapes(slide) -> list:
    return [
        shape
        for shape in iter_shapes(slide.shapes)
        if getattr(shape, "has_text_frame", False)
        and normalize_visible(shape.text) == "Thank you"
    ]


def explicit_font_sizes(shape) -> list[float]:
    values: list[float] = []
    if shape is None or not getattr(shape, "has_text_frame", False):
        return values
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.text.strip() and run.font.size is not None:
                values.append(run.font.size.pt)
    return values


def title_layout_checks(shape, slide_width: int, slide_height: int) -> tuple[bool, dict]:
    if shape is None:
        return False, {"reason": "缺失标题"}
    left, top, width, height = geometry(shape)
    center_x = left + width / 2
    center_y = top + height / 2
    centered_x = abs(center_x - slide_width / 2) <= 1.25 * EMU_PER_INCH
    centered_y = abs(center_y - slide_height / 2) <= 1.50 * EMU_PER_INCH
    large_box = width >= 3.5 * EMU_PER_INCH and height >= 0.55 * EMU_PER_INCH
    sizes = explicit_font_sizes(shape)
    title_placeholder = placeholder_type(shape) in {"TITLE (1)", "CENTER_TITLE (3)"}
    prominent = large_box and (title_placeholder or (sizes and max(sizes) >= 32.0))
    inside = (
        left >= -PAGE_EDGE_TOLERANCE
        and top >= -PAGE_EDGE_TOLERANCE
        and left + width <= slide_width + PAGE_EDGE_TOLERANCE
        and top + height <= slide_height + PAGE_EDGE_TOLERANCE
    )
    passed = centered_x and centered_y and prominent and inside
    return passed, {
        "center_offset_in": (
            round((center_x - slide_width / 2) / EMU_PER_INCH, 4),
            round((center_y - slide_height / 2) / EMU_PER_INCH, 4),
        ),
        "box_in": (
            round(width / EMU_PER_INCH, 4),
            round(height / EMU_PER_INCH, 4),
        ),
        "explicit_font_sizes_pt": sizes,
        "placeholder": placeholder_type(shape),
        "inside": inside,
    }


def extra_final_text(slide, title, slide_height: int) -> list[str]:
    title_element = title._element if title is not None else None
    extras: list[str] = []
    for shape in iter_shapes(slide.shapes):
        if shape._element is title_element or not getattr(shape, "has_text_frame", False):
            continue
        text = normalize_visible(shape.text)
        if not text:
            continue
        if text == "15" and (
            "SLIDE_NUMBER" in placeholder_type(shape)
            or shape.top >= int(0.82 * slide_height)
        ):
            continue
        extras.append(text)
    return extras


def full_slide_new_pictures(slide, slide_width: int, slide_height: int) -> list[str]:
    values: list[str] = []
    for shape in iter_shapes(slide.shapes):
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        coverage = shape.width * shape.height / (slide_width * slide_height)
        if coverage >= 0.85:
            values.append(f"{shape.name}:{coverage:.3f}")
    return values


def render_presentation(path: Path, outdir: Path) -> tuple[list[Path], int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "lo-profile"
    profile.mkdir()
    completed = subprocess.run(
        [
            "soffice",
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(outdir),
            str(path),
        ],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return [], 0, f"转 PDF 失败：{completed.stdout[-500:]}"
    info = subprocess.run(
        ["pdfinfo", str(pdf)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=30,
        check=False,
    ).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    pages = int(match.group(1)) if match else 0
    prefix = outdir / "slide"
    raster = subprocess.run(
        ["pdftoppm", "-jpeg", "-r", "90", str(pdf), str(prefix)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    images = sorted(outdir.glob("slide-*.jpg")) if raster.returncode == 0 else []
    return images, pages, (
        f"页数={pages}，渲染图={len(images)}，"
        f"soffice={completed.returncode}，pdftoppm={raster.returncode}"
    )


def image_similarity(first: Path, second: Path) -> float:
    with Image.open(first) as left_image, Image.open(second) as right_image:
        left = left_image.convert("RGB")
        right = right_image.convert("RGB")
        if left.size != right.size:
            return 0.0
        difference = ImageChops.difference(left, right)
        histogram = difference.histogram()
        squared = sum(count * ((index % 256) ** 2) for index, count in enumerate(histogram))
        pixels = left.width * left.height
        rms = math.sqrt(squared / (pixels * 3)) if pixels else 255.0
        return max(0.0, 1.0 - rms / 255.0)


def emu_rect_to_pixels(
    shape, image_size: tuple[int, int], slide_size: tuple[int, int]
) -> tuple[int, int, int, int]:
    image_width, image_height = image_size
    slide_width, slide_height = slide_size
    left, top, width, height = geometry(shape)
    return (
        max(0, math.floor(left * image_width / slide_width)),
        max(0, math.floor(top * image_height / slide_height)),
        min(image_width, math.ceil((left + width) * image_width / slide_width)),
        min(image_height, math.ceil((top + height) * image_height / slide_height)),
    )


def title_crop_contrast(
    image_path: Path, shape, slide_size: tuple[int, int]
) -> float:
    with Image.open(image_path) as image:
        gray = image.convert("L")
        rectangle = emu_rect_to_pixels(shape, gray.size, slide_size)
        crop = gray.crop(rectangle)
        if crop.width == 0 or crop.height == 0:
            return 0.0
        return float(statistics.mean(ImageStat.Stat(crop).stddev))


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {
        "criterion_id": identifier,
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
        "raw": {
            "description": description,
            "weight": weight,
            "type": kind,
            "passed": bool(passed),
        },
    }


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [
        item["criterion_id"]
        for item in criteria
        if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]
    ]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(
        item["raw"]["weight"]
        for item in criteria
        if item["raw"]["weight"] > 0 and item["raw"]["passed"]
    )
    penalty = sum(
        abs(item["raw"]["weight"])
        for item in criteria
        if item["raw"]["weight"] < 0 and not item["raw"]["passed"]
    )
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [sha256(path)] if path.is_file() else [],
        "criterion_results": criteria,
        "reward": reward,
        "passed": reward >= 0.8 and not hurdles,
        "hurdle_failures": hurdles,
    }


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        criteria = [
            result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配")
            for identifier in SPECS
        ]
        return make_result(path, criteria)
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        criteria = [
            result_item(identifier, False, f"PPTX 打开失败：{exc}")
            for identifier in SPECS
        ]
        return make_result(path, criteria)

    criteria: list[dict] = []
    package_passed, package_evidence = package_ok(path)
    size_passed = (
        len(source.slides) == 14
        and len(candidate.slides) == 15
        and candidate.slide_width == source.slide_width
        and candidate.slide_height == source.slide_height
    )
    criteria.append(
        result_item(
            "R001",
            package_passed and size_passed,
            f"{package_evidence}，输入页数={len(source.slides)}，候选页数={len(candidate.slides)}，"
            f"页面={candidate.slide_width / EMU_PER_INCH:.4f}×"
            f"{candidate.slide_height / EMU_PER_INCH:.4f}英寸",
        )
    )

    content_mismatches: list[int] = []
    style_mismatches: list[int] = []
    notes_mismatches: list[int] = []
    if len(candidate.slides) >= 14:
        for index in range(14):
            source_slide = source.slides[index]
            candidate_slide = candidate.slides[index]
            if slide_content_signature(candidate_slide) != slide_content_signature(source_slide):
                content_mismatches.append(index + 1)
            if slide_style_signature(candidate_slide) != slide_style_signature(source_slide):
                style_mismatches.append(index + 1)
            if slide_notes(candidate_slide) != slide_notes(source_slide):
                notes_mismatches.append(index + 1)
    else:
        content_mismatches = list(range(1, 15))
        style_mismatches = list(range(1, 15))
        notes_mismatches = list(range(1, 15))
    criteria.append(
        result_item(
            "R002",
            not content_mismatches,
            f"原页内容、顺序、图表、图片或结构变化页={content_mismatches}",
        )
    )

    try:
        source_media = media_hashes(source_path)
        candidate_media = media_hashes(path)
        source_media_preserved = all(
            candidate_media[value] >= count for value, count in source_media.items()
        )
    except Exception:
        source_media = Counter()
        candidate_media = Counter()
        source_media_preserved = False
    criteria.append(
        result_item(
            "R003",
            not style_mismatches and not notes_mismatches and source_media_preserved,
            f"原页样式或几何变化页={style_mismatches}，备注变化页={notes_mismatches}，"
            f"原媒体保留={source_media_preserved}，媒体数={sum(source_media.values())}→"
            f"{sum(candidate_media.values())}",
        )
    )

    final_slide = candidate.slides[-1] if len(candidate.slides) else None
    title_matches = thank_you_shapes(final_slide) if final_slide is not None else []
    title = title_matches[0] if len(title_matches) == 1 else None
    editable_title_passed = (
        len(title_matches) == 1
        and title is not None
        and getattr(title, "has_text_frame", False)
        and normalize_visible(title.text) == "Thank you"
    )
    criteria.append(
        result_item(
            "R004",
            editable_title_passed,
            f"末页 Thank you 可编辑文字对象数={len(title_matches)}，"
            f"对象类型={str(title.shape_type) if title is not None else '缺失'}",
        )
    )

    layout_passed, layout_evidence = title_layout_checks(
        title, candidate.slide_width, candidate.slide_height
    )
    criteria.append(
        result_item(
            "R005",
            layout_passed,
            f"标题布局={json.dumps(layout_evidence, ensure_ascii=False)}",
        )
    )

    extras = (
        extra_final_text(final_slide, title, candidate.slide_height)
        if final_slide is not None
        else ["缺失末页"]
    )
    criteria.append(
        result_item("R006", not extras, f"末页无关或重复文字={extras}")
    )

    full_slide_pictures = (
        full_slide_new_pictures(final_slide, candidate.slide_width, candidate.slide_height)
        if final_slide is not None
        else ["缺失末页"]
    )
    native_passed = (
        package_passed
        and editable_title_passed
        and not full_slide_pictures
        and source_media_preserved
    )
    criteria.append(
        result_item(
            "R008",
            native_passed,
            f"末页整页图片={full_slide_pictures}，可编辑标题={editable_title_passed}，"
            f"原媒体保留={source_media_preserved}",
        )
    )

    artifact_files = sorted(
        str(item.relative_to(artifact_dir))
        for item in artifact_dir.rglob("*")
        if item.is_file()
    )
    all_text = "\n".join(
        shape.text
        for slide in candidate.slides
        for shape in iter_shapes(slide.shapes)
        if getattr(shape, "has_text_frame", False)
    )
    all_notes = "\n".join(text for slide in candidate.slides for text in slide_notes(slide))
    lowered = f"{all_text}\n{all_notes}".lower()
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker.lower() in lowered]
    hygiene_passed = artifact_files == [OUTPUT_NAME] and not markers
    criteria.append(
        result_item(
            "R010",
            hygiene_passed,
            f"输出文件={artifact_files}，可疑标记={markers}",
        )
    )

    pre_render_ready = all(
        next(item for item in criteria if item["criterion_id"] == identifier)["raw"]["passed"]
        for identifier in ["R001", "R002", "R003", "R004", "R008"]
    )
    originals_visual_passed = False
    original_visual_min = 0.0
    render_all_passed = False
    title_contrast = 0.0
    render_evidence = "前置结构或内容检查失败，未执行渲染"

    if pre_render_ready:
        with tempfile.TemporaryDirectory(prefix="pptarena023-judge-") as temporary:
            root = Path(temporary)
            source_images, source_pages, source_note = render_presentation(
                source_path, root / "source"
            )
            candidate_images, candidate_pages, candidate_note = render_presentation(
                path, root / "candidate"
            )
            render_evidence = f"源[{source_note}]；候选[{candidate_note}]"
            scores = [
                image_similarity(source_images[index], candidate_images[index])
                for index in range(14)
            ] if len(source_images) == 14 and len(candidate_images) == 15 else []
            original_visual_min = min(scores) if scores else 0.0
            originals_visual_passed = (
                source_pages == 14
                and candidate_pages == 15
                and len(source_images) == 14
                and len(candidate_images) == 15
                and original_visual_min >= 0.9985
            )
            if len(candidate_images) == 15 and title is not None:
                title_contrast = title_crop_contrast(
                    candidate_images[-1],
                    title,
                    (candidate.slide_width, candidate.slide_height),
                )
            render_all_passed = (
                candidate_pages == 15
                and len(candidate_images) == 15
                and title is not None
                and layout_evidence.get("inside", False)
                and title_contrast >= 2.0
            )

    criteria.append(
        result_item(
            "R007",
            originals_visual_passed,
            f"原 14 页最低像素相似度={original_visual_min:.6f}；{render_evidence}",
        )
    )
    criteria.append(
        result_item(
            "R009",
            render_all_passed,
            f"末页标题框灰度标准差={title_contrast:.4f}，"
            f"标题在页内={layout_evidence.get('inside', False)}；{render_evidence}",
        )
    )

    severe = (
        not pre_render_ready
        or not originals_visual_passed
        or not render_all_passed
        or bool(full_slide_pictures)
        or bool(markers)
    )
    criteria.append(result_item("R011", not severe, f"严重错误={severe}"))
    order = {identifier: index for index, identifier in enumerate(SPECS)}
    criteria.sort(key=lambda item: order[item["criterion_id"]])
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    output = args.artifact_dir / OUTPUT_NAME
    if output.is_file():
        evaluation = evaluate(output, args.artifact_dir)
    else:
        evaluation = {
            "task_id": TASK_ID,
            "rubric_version": RUBRIC_VERSION,
            "judge_version": JUDGE_VERSION,
            "artifact_sha256": [],
            "criterion_results": [],
            "reward": 0.0,
            "passed": False,
            "hurdle_failures": ["missing_output"],
        }
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(
        json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
