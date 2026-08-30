#!/usr/bin/env python3
"""PPTArena Case 25 的独立合法解、边界例和负例构造。"""

from __future__ import annotations

import argparse
import io
import shutil
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches


OUTPUT_NAME = "FixTableThemeColor_GroundTruthA.pptx"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS = {"p": P_NS, "a": A_NS}
EMU_PER_POINT = 12_700


def qn(namespace: str, local: str) -> str:
    return f"{{{namespace}}}{local}"


def target(outdir: Path, case: str) -> Path:
    path = outdir / case / OUTPUT_NAME
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def serialize(root) -> bytes:
    return etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)


def rewrite_archive(source: Path, output: Path, transform) -> None:
    with zipfile.ZipFile(source) as archive_in, zipfile.ZipFile(output, "w") as archive_out:
        for info in archive_in.infolist():
            archive_out.writestr(info, transform(info.filename, archive_in.read(info.filename)))


def remove_fill(parent) -> None:
    names = {"noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill"}
    for child in list(parent):
        if isinstance(child.tag, str) and etree.QName(child).localname in names:
            parent.remove(child)


def set_fill(parent, spec: tuple[str, str, tuple[tuple[str, str], ...]]) -> None:
    kind, value, transforms = spec
    remove_fill(parent)
    fill = etree.Element(qn(A_NS, "solidFill"))
    color = etree.SubElement(fill, qn(A_NS, kind), val=value)
    for name, amount in transforms:
        etree.SubElement(color, qn(A_NS, name), val=amount)
    parent.insert(0, fill)


def scheme(value: str, *transforms: tuple[str, str]) -> tuple[str, str, tuple[tuple[str, str], ...]]:
    return "schemeClr", value, tuple(transforms)


def srgb(value: str) -> tuple[str, str, tuple[tuple[str, str], ...]]:
    return "srgbClr", value, ()


HEADER = scheme("accent6", ("lumMod", "50000"))
ROW_A = scheme("tx2", ("lumMod", "75000"), ("lumOff", "25000"))
ROW_B = scheme("tx2", ("lumMod", "50000"), ("lumOff", "50000"))
WHITE = scheme("bg1")
BLACK = scheme("dk1")


def style_variant(
    source: Path,
    output: Path,
    row_fills: list[tuple[str, str, tuple[tuple[str, str], ...]]],
    row_text: list[tuple[str, str, tuple[tuple[str, str], ...]]],
    cell_override: dict[tuple[int, int], tuple[tuple[str, str, tuple[tuple[str, str], ...]], tuple[str, str, tuple[tuple[str, str], ...]]]] | None = None,
) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name != "ppt/slides/slide1.xml":
            return data
        root = etree.fromstring(data)
        tables = root.xpath("./p:cSld/p:spTree/p:graphicFrame/a:graphic/a:graphicData/a:tbl", namespaces=NS)
        if len(tables) != 1:
            raise ValueError("表格不唯一")
        rows = tables[0].xpath("./a:tr", namespaces=NS)
        for row_index, row in enumerate(rows):
            for column_index, cell in enumerate(row.xpath("./a:tc", namespaces=NS)):
                fill_spec = row_fills[row_index]
                text_spec = row_text[row_index]
                if cell_override and (row_index, column_index) in cell_override:
                    fill_spec, text_spec = cell_override[(row_index, column_index)]
                properties = cell.find("a:tcPr", NS)
                set_fill(properties, fill_spec)
                for text_properties in cell.xpath("./a:txBody/a:p/a:r/a:rPr | ./a:txBody/a:p/a:endParaRPr", namespaces=NS):
                    set_fill(text_properties, text_spec)
        return serialize(root)

    rewrite_archive(source, output, transform)


def mutate_slide(source: Path, output: Path, mutate) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name != "ppt/slides/slide1.xml":
            return data
        root = etree.fromstring(data)
        mutate(root)
        return serialize(root)

    rewrite_archive(source, output, transform)


def table_frame(root):
    nodes = root.xpath("./p:cSld/p:spTree/p:graphicFrame[a:graphic/a:graphicData/a:tbl]", namespaces=NS)
    if len(nodes) != 1:
        raise ValueError("表格框不唯一")
    return nodes[0]


def change_text(source: Path, output: Path, old: str, new: str) -> None:
    def mutate(root) -> None:
        matches = [node for node in root.xpath(".//a:t", namespaces=NS) if node.text == old]
        if len(matches) != 1:
            raise ValueError(f"文字 {old!r} 不唯一")
        matches[0].text = new

    mutate_slide(source, output, mutate)


def move_or_resize_table(source: Path, output: Path, *, left_delta_pt: float = 0.0, width_delta_pt: float = 0.0) -> None:
    def mutate(root) -> None:
        transform = table_frame(root).find("./p:xfrm", NS)
        offset = transform.find("a:off", NS)
        extent = transform.find("a:ext", NS)
        offset.set("x", str(int(offset.get("x")) + round(left_delta_pt * EMU_PER_POINT)))
        extent.set("cx", str(int(extent.get("cx")) + round(width_delta_pt * EMU_PER_POINT)))

    mutate_slide(source, output, mutate)


def change_row_height(source: Path, output: Path) -> None:
    def mutate(root) -> None:
        row = table_frame(root).xpath("./a:graphic/a:graphicData/a:tbl/a:tr", namespaces=NS)[1]
        row.set("h", str(int(row.get("h")) + 12 * EMU_PER_POINT))

    mutate_slide(source, output, mutate)


def change_column_width(source: Path, output: Path) -> None:
    def mutate(root) -> None:
        columns = table_frame(root).xpath("./a:graphic/a:graphicData/a:tbl/a:tblGrid/a:gridCol", namespaces=NS)
        columns[0].set("w", str(int(columns[0].get("w")) + 12 * EMU_PER_POINT))
        columns[1].set("w", str(int(columns[1].get("w")) - 12 * EMU_PER_POINT))

    mutate_slide(source, output, mutate)


def remove_table(source: Path, output: Path) -> None:
    def mutate(root) -> None:
        frame = table_frame(root)
        frame.getparent().remove(frame)

    mutate_slide(source, output, mutate)


def repack_content_types(source: Path, output: Path) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name != "[Content_Types].xml":
            return data
        root = etree.fromstring(data)
        root.append(etree.Comment("harmless repack"))
        return serialize(root)

    rewrite_archive(source, output, transform)


def change_theme(source: Path, output: Path) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name != "ppt/theme/theme1.xml":
            return data
        root = etree.fromstring(data)
        accent = root.find("./a:themeElements/a:clrScheme/a:accent6/a:srgbClr", NS)
        accent.set("val", "AA66CC")
        return serialize(root)

    rewrite_archive(source, output, transform)


def change_media(source: Path, output: Path) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name != "ppt/media/image1.png":
            return data
        image = Image.open(io.BytesIO(data)).convert("RGBA")
        pixel = image.getpixel((0, 0))
        image.putpixel((0, 0), ((pixel[0] + 1) % 256, pixel[1], pixel[2], pixel[3]))
        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        return buffer.getvalue()

    rewrite_archive(source, output, transform)


def add_textbox(source: Path, output: Path, text: str) -> None:
    presentation = Presentation(source)
    box = presentation.slides[0].shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(3.0), Inches(0.35))
    box.text = text
    presentation.save(output)


def add_slide(source: Path, output: Path) -> None:
    presentation = Presentation(source)
    presentation.slides.add_slide(presentation.slide_layouts[0])
    presentation.save(output)


def flatten_table(source: Path, output: Path) -> None:
    presentation = Presentation(source)
    slide = presentation.slides[0]
    table = next(shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.TABLE)
    left, top, width, height = table.left, table.top, table.width, table.height
    table.element.getparent().remove(table.element)
    with tempfile.TemporaryDirectory(prefix="pptarena025-flat-") as temporary:
        image_path = Path(temporary) / "flat.png"
        image = Image.new("RGB", (1200, 460), (68, 105, 100))
        image.save(image_path)
        slide.shapes.add_picture(str(image_path), left, top, width, height)
        presentation.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--upstream-ground-truth", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()

    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    target(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, target(args.outdir, "unmodified"))
    shutil.copy2(args.upstream_ground_truth, target(args.outdir, "upstream-ground-truth-media-mutated"))

    style_variant(
        args.source,
        target(args.outdir, "valid-accent3"),
        [
            scheme("accent3", ("lumMod", "45000")),
            scheme("accent3", ("lumMod", "70000"), ("lumOff", "30000")),
            scheme("accent3", ("lumMod", "50000"), ("lumOff", "50000")),
            scheme("accent3", ("lumMod", "70000"), ("lumOff", "30000")),
            scheme("accent3", ("lumMod", "50000"), ("lumOff", "50000")),
        ],
        [WHITE, BLACK, BLACK, BLACK, BLACK],
    )
    style_variant(
        args.source,
        target(args.outdir, "valid-dk2-accent6"),
        [
            scheme("dk2"),
            scheme("accent6", ("lumMod", "80000"), ("lumOff", "20000")),
            scheme("accent6", ("lumMod", "60000"), ("lumOff", "40000")),
            scheme("accent6", ("lumMod", "80000"), ("lumOff", "20000")),
            scheme("accent6", ("lumMod", "60000"), ("lumOff", "40000")),
        ],
        [WHITE, BLACK, BLACK, BLACK, BLACK],
    )
    repack_content_types(args.golden, target(args.outdir, "repacked-content-types"))

    style_variant(args.source, target(args.outdir, "light-header"), [scheme("accent6", ("lumMod", "50000"), ("lumOff", "50000")), ROW_A, ROW_B, ROW_A, ROW_B], [BLACK, WHITE, WHITE, WHITE, WHITE])
    style_variant(args.source, target(args.outdir, "non-theme-header"), [srgb("242060"), ROW_A, ROW_B, ROW_A, ROW_B], [WHITE, WHITE, WHITE, WHITE, WHITE])
    style_variant(args.source, target(args.outdir, "dark-header-text"), [HEADER, ROW_A, ROW_B, ROW_A, ROW_B], [BLACK, WHITE, WHITE, WHITE, WHITE])
    style_variant(args.source, target(args.outdir, "no-banding"), [HEADER, ROW_A, ROW_A, ROW_A, ROW_A], [WHITE, WHITE, WHITE, WHITE, WHITE])
    style_variant(args.source, target(args.outdir, "wrong-row-sequence"), [HEADER, ROW_A, ROW_A, ROW_B, ROW_B], [WHITE, WHITE, WHITE, WHITE, WHITE])
    style_variant(args.source, target(args.outdir, "inconsistent-row"), [HEADER, ROW_A, ROW_B, ROW_A, ROW_B], [WHITE, WHITE, WHITE, WHITE, WHITE], {(2, 3): (ROW_A, WHITE)})
    style_variant(
        args.source,
        target(args.outdir, "bands-too-similar"),
        [HEADER, scheme("accent6", ("lumMod", "70000"), ("lumOff", "30000")), scheme("accent6", ("lumMod", "68000"), ("lumOff", "32000")), scheme("accent6", ("lumMod", "70000"), ("lumOff", "30000")), scheme("accent6", ("lumMod", "68000"), ("lumOff", "32000"))],
        [WHITE, BLACK, BLACK, BLACK, BLACK],
    )
    style_variant(args.source, target(args.outdir, "yellow-residue"), [HEADER, ROW_A, ROW_B, ROW_A, ROW_B], [WHITE, WHITE, WHITE, WHITE, WHITE], {(4, 2): (srgb("FFFF00"), BLACK)})
    style_variant(args.source, target(args.outdir, "unreadable-data-text"), [HEADER, ROW_A, ROW_B, ROW_A, ROW_B], [WHITE, ROW_A, ROW_B, ROW_A, ROW_B])
    style_variant(args.source, target(args.outdir, "generic-white-gray"), [scheme("dk2"), scheme("lt1"), scheme("lt2"), scheme("lt1"), scheme("lt2")], [WHITE, BLACK, BLACK, BLACK, BLACK])

    change_text(args.golden, target(args.outdir, "changed-table-value"), "Q1", "Q5")
    move_or_resize_table(args.golden, target(args.outdir, "moved-table"), left_delta_pt=10.0)
    move_or_resize_table(args.golden, target(args.outdir, "resized-table"), width_delta_pt=12.0)
    change_row_height(args.golden, target(args.outdir, "changed-row-height"))
    change_column_width(args.golden, target(args.outdir, "changed-column-width"))
    change_text(args.golden, target(args.outdir, "changed-title"), "Financials", "Financial Results")
    change_text(args.golden, target(args.outdir, "changed-slide-number"), "1", "2")
    change_media(args.golden, target(args.outdir, "changed-media"))
    change_theme(args.golden, target(args.outdir, "changed-theme"))
    remove_table(args.golden, target(args.outdir, "missing-table"))
    flatten_table(args.golden, target(args.outdir, "flattened-table"))
    add_textbox(args.golden, target(args.outdir, "extra-object"), "extra object")
    add_textbox(args.golden, target(args.outdir, "hardcoded-marker"), "solution/Golden")
    add_slide(args.golden, target(args.outdir, "extra-slide"))


if __name__ == "__main__":
    main()
