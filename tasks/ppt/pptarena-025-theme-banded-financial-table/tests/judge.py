#!/usr/bin/env python3
"""PPTArena Case 25 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import colorsys
import copy
import hashlib
import json
import math
import re
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from lxml import etree
from PIL import Image, ImageChops
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "FixTableThemeColor_GroundTruthA.pptx"
INPUT_NAME = "FixTableThemeColor_TestA.pptx"
INPUT_SHA256 = "d0b3c9d5b9988a375252315d3300b3237f25684af23ed77a4330aa8021d10c30"
TASK_ID = "cowork-evo/pptarena-025-theme-banded-financial-table"
RUBRIC_VERSION = "pptarena-case-025-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-025-deterministic-v0.1"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS = {"p": P_NS, "a": A_NS}
EMU_PER_POINT = 12_700
EXPECTED_TEXT = [
    ["Quarter", "Revenue ($)", "Expenses ($)", "Net profit ($)"],
    ["Q1", "$200,000", "$150,000", "$50,000"],
    ["Q2", "$300,000", "$200,000", "$100,000"],
    ["Q3", "$400,000", "$250,000", "$150,000"],
    ["Q4", "$300,000", "$200,000", "$100,000"],
]
THEME_FILL_SCHEMES = {"dk2", "tx2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6"}
SUSPICIOUS_MARKERS = [
    "solution/golden",
    "/root/",
    "judge-result",
    "artifact_dir",
    "pptarena-case-025",
]
SPECS = {
    "R001": ("指定 PPTX 可打开、页面尺寸不变且只有 1 页", 6, "hurdle"),
    "R002": ("原生 5×4 表格的内容、结构和几何保持", 17, "hurdle"),
    "R003": ("表头为统一的深色主题填充和高对比浅色文字", 18, "hurdle"),
    "R004": ("数据行为行内同色的 A-B-A-B 行带并保持可读", 18, "hurdle"),
    "R005": ("填充可追溯到原主题色且无亮黄残留", 14, "hurdle"),
    "R006": ("页面非目标内容、主题、母版、版式和媒体保持", 14, "hurdle"),
    "R007": ("渲染完整，表外区域保持且表内可读", 6, "hurdle"),
    "R008": ("保留原生表格和原 3 个对象，无压平或额外对象", 4, "hurdle"),
    "R009": ("输出目录干净且无答案标记", 2, "positive"),
    "R010": ("规范化允许的表格样式后关键内容保持", 1, "positive"),
    "R011": ("严重表格、可读性、可编辑性或保持错误", -20, "penalty"),
}


def qn(namespace: str, local: str) -> str:
    return f"{{{namespace}}}{local}"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def locate_source() -> Path | None:
    candidates = [
        Path("/app/input_files") / INPUT_NAME,
        Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME,
    ]
    return next((path for path in candidates if path.is_file()), None)


def package_ok(path: Path) -> tuple[bool, str]:
    try:
        with zipfile.ZipFile(path) as archive:
            broken = archive.testzip()
            names = {name for name in archive.namelist() if not name.endswith("/")}
            required = {
                "[Content_Types].xml",
                "ppt/presentation.xml",
                "ppt/slides/slide1.xml",
                "ppt/theme/theme1.xml",
            }
            missing = sorted(required - names)
            macros = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
        return broken is None and not missing and not macros, f"损坏成员={broken}，缺失={missing}，宏={macros}"
    except Exception as exc:
        return False, f"ZIP 检查失败：{exc}"


def media_hashes(path: Path) -> Counter[str]:
    with zipfile.ZipFile(path) as archive:
        return Counter(
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("ppt/media/") and not name.endswith("/")
        )


def member_names(path: Path) -> set[str]:
    with zipfile.ZipFile(path) as archive:
        return {name for name in archive.namelist() if not name.endswith("/")}


def theme_colors(path: Path) -> dict[str, tuple[int, int, int]]:
    with zipfile.ZipFile(path) as archive:
        root = etree.fromstring(archive.read("ppt/theme/theme1.xml"))
    colors: dict[str, tuple[int, int, int]] = {}
    nodes = root.xpath("./a:themeElements/a:clrScheme/*", namespaces=NS)
    for node in nodes:
        key = etree.QName(node).localname
        srgb = node.find("a:srgbClr", NS)
        system = node.find("a:sysClr", NS)
        value = srgb.get("val") if srgb is not None else system.get("lastClr") if system is not None else None
        if value and re.fullmatch(r"[0-9A-Fa-f]{6}", value):
            colors[key] = tuple(int(value[index : index + 2], 16) for index in (0, 2, 4))
    return colors


def scheme_key(value: str) -> str:
    return {
        "bg1": "lt1",
        "tx1": "dk1",
        "bg2": "lt2",
        "tx2": "dk2",
    }.get(value, value)


def apply_color_transforms(rgb: tuple[int, int, int], node) -> tuple[int, int, int]:
    red, green, blue = (channel / 255.0 for channel in rgb)
    hue, lightness, saturation = colorsys.rgb_to_hls(red, green, blue)
    for transform in node:
        name = etree.QName(transform).localname
        try:
            value = int(transform.get("val", "100000")) / 100000.0
        except ValueError:
            continue
        if name == "lumMod":
            lightness *= value
        elif name == "lumOff":
            lightness += value
        elif name == "tint":
            lightness += (1.0 - lightness) * value
        elif name == "shade":
            lightness *= value
        elif name == "satMod":
            saturation *= value
        elif name == "satOff":
            saturation += value
    red, green, blue = colorsys.hls_to_rgb(hue, max(0.0, min(1.0, lightness)), max(0.0, min(1.0, saturation)))
    return tuple(round(channel * 255) for channel in (red, green, blue))


def color_node_info(fill, colors: dict[str, tuple[int, int, int]]) -> dict:
    if fill is None:
        return {"kind": "none", "value": None, "rgb": None, "theme_fill": False}
    color_nodes = [child for child in fill if isinstance(child.tag, str)]
    if not color_nodes:
        return {"kind": "none", "value": None, "rgb": None, "theme_fill": False}
    node = color_nodes[0]
    kind = etree.QName(node).localname
    value = node.get("val", "")
    rgb = None
    theme_fill = False
    if kind == "schemeClr":
        key = scheme_key(value)
        if key in colors:
            rgb = apply_color_transforms(colors[key], node)
        theme_fill = value in THEME_FILL_SCHEMES or key in THEME_FILL_SCHEMES
    elif kind == "srgbClr" and re.fullmatch(r"[0-9A-Fa-f]{6}", value):
        rgb = tuple(int(value[index : index + 2], 16) for index in (0, 2, 4))
        rgb = apply_color_transforms(rgb, node)
    return {"kind": kind, "value": value, "rgb": rgb, "theme_fill": theme_fill}


def theme_palette(colors: dict[str, tuple[int, int, int]]) -> list[tuple[int, int, int]]:
    values: list[tuple[int, int, int]] = []
    for name in sorted(THEME_FILL_SCHEMES):
        base = colors.get(scheme_key(name))
        if base is None:
            continue
        red, green, blue = (channel / 255.0 for channel in base)
        hue, lightness, saturation = colorsys.rgb_to_hls(red, green, blue)
        for step in range(0, 19):
            amount = step / 20.0
            for transformed_lightness in (lightness * (1.0 - amount), lightness + (1.0 - lightness) * amount):
                transformed = colorsys.hls_to_rgb(hue, max(0.0, min(1.0, transformed_lightness)), saturation)
                values.append(tuple(round(channel * 255) for channel in transformed))
    return values


def color_distance(first: tuple[int, int, int], second: tuple[int, int, int]) -> float:
    return math.sqrt(sum((a - b) ** 2 for a, b in zip(first, second)))


def is_theme_related(info: dict, rgb: tuple[int, int, int] | None, palette: list[tuple[int, int, int]]) -> bool:
    if info.get("kind") == "schemeClr":
        return bool(info.get("theme_fill"))
    actual = info.get("rgb") or rgb
    return actual is not None and bool(palette) and min(color_distance(actual, item) for item in palette) <= 14.0


def relative_luminance(rgb: tuple[int, int, int]) -> float:
    def linear(channel: int) -> float:
        value = channel / 255.0
        return value / 12.92 if value <= 0.04045 else ((value + 0.055) / 1.055) ** 2.4

    red, green, blue = (linear(channel) for channel in rgb)
    return 0.2126 * red + 0.7152 * green + 0.0722 * blue


def contrast_ratio(first: tuple[int, int, int], second: tuple[int, int, int]) -> float:
    one, two = relative_luminance(first), relative_luminance(second)
    return (max(one, two) + 0.05) / (min(one, two) + 0.05)


def table_xml(path: Path):
    with zipfile.ZipFile(path) as archive:
        root = etree.fromstring(archive.read("ppt/slides/slide1.xml"))
    tables = root.xpath("./p:cSld/p:spTree/p:graphicFrame/a:graphic/a:graphicData/a:tbl", namespaces=NS)
    if len(tables) != 1:
        raise ValueError(f"原生表格数={len(tables)}")
    rows = tables[0].xpath("./a:tr", namespaces=NS)
    return root, tables[0], rows


def xml_cell_styles(path: Path, colors: dict[str, tuple[int, int, int]]) -> list[list[dict]]:
    _, _, rows = table_xml(path)
    values = []
    for row in rows:
        current = []
        for cell in row.xpath("./a:tc", namespaces=NS):
            properties = cell.find("a:tcPr", NS)
            fill = properties.find("a:solidFill", NS) if properties is not None else None
            text_infos = []
            for properties_node in cell.xpath("./a:txBody/a:p/a:r/a:rPr | ./a:txBody/a:p/a:endParaRPr", namespaces=NS):
                text_infos.append(color_node_info(properties_node.find("a:solidFill", NS), colors))
            current.append({"fill": color_node_info(fill, colors), "text": text_infos})
        values.append(current)
    return values


def table_shape(presentation: Presentation):
    if len(presentation.slides) != 1:
        return None
    tables = [shape for shape in presentation.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.TABLE]
    return tables[0] if len(tables) == 1 else None


def table_inventory(presentation: Presentation) -> dict:
    shape = table_shape(presentation)
    if shape is None:
        return {"passed": False, "reason": "原生表格不唯一"}
    table = shape.table
    text = [[table.cell(row, column).text for column in range(len(table.columns))] for row in range(len(table.rows))]
    merges = [
        [
            [bool(table.cell(row, column).is_merge_origin), bool(table.cell(row, column).is_spanned)]
            for column in range(len(table.columns))
        ]
        for row in range(len(table.rows))
    ]
    cells = []
    for row in range(len(table.rows)):
        current = []
        for column in range(len(table.columns)):
            cell = table.cell(row, column)
            paragraphs = []
            for paragraph in cell.text_frame.paragraphs:
                paragraphs.append({
                    "alignment": str(paragraph.alignment),
                    "level": paragraph.level,
                    "text": paragraph.text,
                })
            current.append({
                "margin_left": int(cell.margin_left),
                "margin_right": int(cell.margin_right),
                "margin_top": int(cell.margin_top),
                "margin_bottom": int(cell.margin_bottom),
                "vertical_anchor": str(cell.vertical_anchor),
                "paragraphs": paragraphs,
            })
        cells.append(current)
    return {
        "passed": True,
        "text": text,
        "rows": len(table.rows),
        "columns": len(table.columns),
        "left": int(shape.left),
        "top": int(shape.top),
        "width": int(shape.width),
        "height": int(shape.height),
        "row_heights": [int(row.height) for row in table.rows],
        "column_widths": [int(column.width) for column in table.columns],
        "merges": merges,
        "cells": cells,
    }


def shape_c14n(shape) -> str:
    return hashlib.sha256(etree.tostring(shape.element, method="c14n")).hexdigest()


def non_table_signatures(presentation: Presentation) -> list[tuple]:
    if len(presentation.slides) != 1:
        return []
    return [
        (index, int(shape.shape_id), shape.name, str(shape.shape_type), shape_c14n(shape))
        for index, shape in enumerate(presentation.slides[0].shapes)
        if shape.shape_type != MSO_SHAPE_TYPE.TABLE
    ]


def notes_signature(presentation: Presentation) -> tuple[str, ...]:
    values = []
    for slide in presentation.slides:
        try:
            values.append(slide.notes_slide.notes_text_frame.text)
        except Exception:
            values.append("")
    return tuple(values)


def critical_hashes(path: Path) -> dict[str, str]:
    prefixes = (
        "ppt/theme/",
        "ppt/slideMasters/",
        "ppt/slideLayouts/",
        "ppt/notesMasters/",
        "ppt/notesSlides/",
        "ppt/media/",
        "ppt/slides/_rels/",
    )
    with zipfile.ZipFile(path) as archive:
        return {
            name: hashlib.sha256(archive.read(name)).hexdigest()
            for name in sorted(archive.namelist())
            if not name.endswith("/") and name.startswith(prefixes)
        }


def scaffold_hashes(path: Path) -> dict[str, str]:
    parts = {
        "[Content_Types].xml",
        "_rels/.rels",
        "ppt/presentation.xml",
        "ppt/_rels/presentation.xml.rels",
    }
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        return {
            name: hashlib.sha256(archive.read(name)).hexdigest()
            for name in sorted(parts & names)
        }


def remove_fill(parent) -> None:
    fill_names = {"noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill"}
    for child in list(parent):
        if isinstance(child.tag, str) and etree.QName(child).localname in fill_names:
            parent.remove(child)


def normalized_slide_hash(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        root = etree.fromstring(archive.read("ppt/slides/slide1.xml"))
    tables = root.xpath("./p:cSld/p:spTree/p:graphicFrame/a:graphic/a:graphicData/a:tbl", namespaces=NS)
    if len(tables) != 1:
        raise ValueError("无法规范化表格")
    table = tables[0]
    table_properties = table.find("a:tblPr", NS)
    if table_properties is not None:
        table_properties.attrib.pop("firstRow", None)
        table_properties.attrib.pop("bandRow", None)
        for style in list(table_properties.findall("a:tableStyleId", NS)):
            table_properties.remove(style)
    for cell_properties in table.xpath("./a:tr/a:tc/a:tcPr", namespaces=NS):
        remove_fill(cell_properties)
    for text_properties in table.xpath(".//a:rPr | .//a:endParaRPr", namespaces=NS):
        remove_fill(text_properties)
    common_slide = root.find(qn(P_NS, "cSld"))
    return hashlib.sha256(etree.tostring(common_slide, method="c14n")).hexdigest()


def render_presentation(path: Path, outdir: Path) -> tuple[list[Path], int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "lo-profile"
    profile.mkdir()
    completed = subprocess.run(
        [
            "soffice",
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(outdir),
            str(path),
        ],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return [], 0, f"转 PDF 失败：{completed.stdout[-500:]}"
    info = subprocess.run(
        ["pdfinfo", str(pdf)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=30,
        check=False,
    ).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    pages = int(match.group(1)) if match else 0
    raster = subprocess.run(
        ["pdftoppm", "-jpeg", "-r", "150", str(pdf), str(outdir / "slide")],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    images = sorted(outdir.glob("slide-*.jpg")) if raster.returncode == 0 else []
    return images, pages, f"页数={pages}，渲染图={len(images)}，soffice={completed.returncode}，pdftoppm={raster.returncode}"


def table_pixel_bounds(image: Image.Image, inventory: dict, slide_size: tuple[int, int]) -> tuple[int, int, int, int]:
    left = round(inventory["left"] * image.width / slide_size[0])
    top = round(inventory["top"] * image.height / slide_size[1])
    right = round((inventory["left"] + inventory["width"]) * image.width / slide_size[0])
    bottom = round((inventory["top"] + inventory["height"]) * image.height / slide_size[1])
    return left, top, right, bottom


def outside_table_similarity(source_image: Image.Image, candidate_image: Image.Image, inventory: dict, slide_size: tuple[int, int]) -> tuple[float, dict]:
    source = source_image.convert("RGB")
    candidate = candidate_image.convert("RGB")
    if source.size != candidate.size:
        return 0.0, {"reason": "渲染尺寸不一致"}
    left, top, right, bottom = table_pixel_bounds(source, inventory, slide_size)
    padding = max(8, round(min(source.size) * 0.008))
    mask = Image.new("L", source.size, 255)
    mask.paste(0, (max(0, left - padding), max(0, top - padding), min(source.width, right + padding), min(source.height, bottom + padding)))
    difference = ImageChops.difference(source, candidate)
    pixels = mask.histogram()[255]
    histogram = difference.histogram(mask)
    squared = sum(count * ((index % 256) ** 2) for index, count in enumerate(histogram))
    rms = math.sqrt(squared / (pixels * 3)) if pixels else 255.0
    return max(0.0, 1.0 - rms / 255.0), {
        "table_pixels": [left, top, right, bottom],
        "padding": padding,
        "evaluated_pixels": pixels,
        "rms": round(rms, 6),
    }


def dominant_color(image: Image.Image, bounds: tuple[int, int, int, int]) -> tuple[int, int, int]:
    crop = image.convert("RGB").crop(bounds)
    counter: Counter[tuple[int, int, int]] = Counter()
    for red, green, blue in crop.getdata():
        counter[(round(red / 4) * 4, round(green / 4) * 4, round(blue / 4) * 4)] += 1
    return counter.most_common(1)[0][0]


def rendered_cell_colors(image: Image.Image, inventory: dict, slide_size: tuple[int, int]) -> list[list[tuple[int, int, int]]]:
    table_left, table_top, _, _ = table_pixel_bounds(image, inventory, slide_size)
    x_values = [table_left]
    y_values = [table_top]
    for width in inventory["column_widths"]:
        x_values.append(round(x_values[-1] + width * image.width / slide_size[0]))
    for height in inventory["row_heights"]:
        y_values.append(round(y_values[-1] + height * image.height / slide_size[1]))
    values = []
    inset = max(12, round(min(image.size) * 0.008))
    for row in range(5):
        current = []
        for column in range(4):
            bounds = (
                x_values[column] + inset,
                y_values[row] + inset,
                x_values[column + 1] - inset,
                y_values[row + 1] - inset,
            )
            current.append(dominant_color(image, bounds))
        values.append(current)
    return values


def resolved_text_colors(cell: dict) -> list[tuple[int, int, int]]:
    values = [item["rgb"] for item in cell["text"] if item.get("rgb") is not None]
    return values


def row_uniform(colors: list[tuple[int, int, int]], tolerance: float = 12.0) -> bool:
    return bool(colors) and max(color_distance(first, second) for first in colors for second in colors) <= tolerance


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {
        "criterion_id": identifier,
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
        "raw": {
            "description": description,
            "weight": weight,
            "type": kind,
            "passed": bool(passed),
        },
    }


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [
        item["criterion_id"]
        for item in criteria
        if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]
    ]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(
        item["raw"]["weight"]
        for item in criteria
        if item["raw"]["weight"] > 0 and item["raw"]["passed"]
    )
    penalty = sum(
        abs(item["raw"]["weight"])
        for item in criteria
        if item["raw"]["weight"] < 0 and not item["raw"]["passed"]
    )
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [sha256(path)] if path.is_file() else [],
        "criterion_results": criteria,
        "reward": reward,
        "passed": reward >= 0.8 and not hurdles,
        "hurdle_failures": hurdles,
    }


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        return make_result(path, [result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配") for identifier in SPECS])
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        return make_result(path, [result_item(identifier, False, f"PPTX 打开失败：{exc}") for identifier in SPECS])

    criteria: list[dict] = []
    package_passed, package_evidence = package_ok(path)
    size_passed = (
        len(source.slides) == 1
        and len(candidate.slides) == 1
        and candidate.slide_width == source.slide_width
        and candidate.slide_height == source.slide_height
    )
    criteria.append(result_item("R001", package_passed and size_passed, f"{package_evidence}，页数={len(candidate.slides)}，页面={candidate.slide_width}×{candidate.slide_height}"))

    source_inventory = table_inventory(source)
    candidate_inventory = table_inventory(candidate)
    structure_keys = ["text", "rows", "columns", "left", "top", "width", "height", "row_heights", "column_widths", "merges", "cells"]
    table_passed = (
        source_inventory.get("passed")
        and candidate_inventory.get("passed")
        and candidate_inventory.get("text") == EXPECTED_TEXT
        and all(candidate_inventory.get(key) == source_inventory.get(key) for key in structure_keys)
    )
    criteria.append(result_item("R002", bool(table_passed), json.dumps({key: candidate_inventory.get(key) for key in structure_keys[:-1]}, ensure_ascii=False)))

    try:
        colors = theme_colors(source_path)
        palette = theme_palette(colors)
        style_xml = xml_cell_styles(path, colors)
    except Exception as exc:
        colors, palette, style_xml = {}, [], []
        style_error = str(exc)
    else:
        style_error = ""

    render_note = "前置检查失败，未执行渲染"
    rendered_colors: list[list[tuple[int, int, int]]] = []
    similarity = 0.0
    similarity_evidence: dict = {}
    render_passed = False
    if package_passed and size_passed and candidate_inventory.get("passed"):
        with tempfile.TemporaryDirectory(prefix="pptarena025-judge-") as temporary:
            root = Path(temporary)
            source_images, source_pages, source_note = render_presentation(source_path, root / "source")
            candidate_images, candidate_pages, candidate_note = render_presentation(path, root / "candidate")
            render_note = f"源[{source_note}]；候选[{candidate_note}]"
            if len(source_images) == 1 and len(candidate_images) == 1 and source_pages == candidate_pages == 1:
                with Image.open(source_images[0]) as source_image, Image.open(candidate_images[0]) as candidate_image:
                    similarity, similarity_evidence = outside_table_similarity(
                        source_image,
                        candidate_image,
                        source_inventory,
                        (source.slide_width, source.slide_height),
                    )
                    rendered_colors = rendered_cell_colors(
                        candidate_image,
                        candidate_inventory,
                        (candidate.slide_width, candidate.slide_height),
                    )
                render_passed = similarity >= 0.9985 and len(rendered_colors) == 5

    header_passed = False
    banding_passed = False
    theme_passed = False
    header_details: dict = {}
    banding_details: dict = {}
    theme_details: dict = {}
    if len(rendered_colors) == 5 and len(style_xml) == 5:
        header_fills = rendered_colors[0]
        header_text_values = [resolved_text_colors(cell) for cell in style_xml[0]]
        header_contrasts = []
        header_light = []
        for fill, texts in zip(header_fills, header_text_values):
            if texts:
                header_contrasts.append(min(contrast_ratio(fill, text) for text in texts))
                header_light.append(min(relative_luminance(text) for text in texts))
            else:
                header_contrasts.append(0.0)
                header_light.append(0.0)
        header_passed = (
            row_uniform(header_fills)
            and max(relative_luminance(fill) for fill in header_fills) <= 0.24
            and min(header_contrasts) >= 4.5
            and min(header_light) >= 0.72
        )
        header_details = {
            "fills": header_fills,
            "luminance": [round(relative_luminance(fill), 4) for fill in header_fills],
            "text_contrast": [round(value, 4) for value in header_contrasts],
            "text_luminance": [round(value, 4) for value in header_light],
        }

        row_fills = rendered_colors[1:]
        row_colors = [row[0] for row in row_fills]
        data_contrasts = []
        for row_index, row in enumerate(row_fills, start=1):
            current = []
            for fill, cell in zip(row, style_xml[row_index]):
                texts = resolved_text_colors(cell)
                current.append(min((contrast_ratio(fill, text) for text in texts), default=0.0))
            data_contrasts.append(current)
        banding_passed = (
            all(row_uniform(row) for row in row_fills)
            and color_distance(row_colors[0], row_colors[2]) <= 14.0
            and color_distance(row_colors[1], row_colors[3]) <= 14.0
            and color_distance(row_colors[0], row_colors[1]) >= 28.0
            and min(value for row in data_contrasts for value in row) >= 2.55
        )
        banding_details = {
            "row_fills": row_fills,
            "AB_distance": round(color_distance(row_colors[0], row_colors[1]), 4),
            "AA_distance": round(color_distance(row_colors[0], row_colors[2]), 4),
            "BB_distance": round(color_distance(row_colors[1], row_colors[3]), 4),
            "text_contrast": [[round(value, 4) for value in row] for row in data_contrasts],
        }

        theme_checks = []
        yellow_distances = []
        for row_index, row in enumerate(style_xml):
            for column_index, cell in enumerate(row):
                rendered = rendered_colors[row_index][column_index]
                theme_checks.append(is_theme_related(cell["fill"], rendered, palette))
                yellow_distances.append(color_distance(rendered, (255, 255, 0)))
        theme_passed = all(theme_checks) and min(yellow_distances) >= 90.0
        theme_details = {
            "theme_checks": theme_checks,
            "minimum_yellow_distance": round(min(yellow_distances), 4),
            "xml_fills": [[cell["fill"] for cell in row] for row in style_xml],
        }
    criteria.append(result_item("R003", header_passed, f"{json.dumps(header_details, ensure_ascii=False)}；样式解析={style_error}"))
    criteria.append(result_item("R004", banding_passed, json.dumps(banding_details, ensure_ascii=False)))
    criteria.append(result_item("R005", theme_passed, json.dumps(theme_details, ensure_ascii=False)))

    try:
        media_equal = media_hashes(path) == media_hashes(source_path)
        media_count = sum(media_hashes(path).values())
        critical_equal = critical_hashes(path) == critical_hashes(source_path)
        members_equal = member_names(path) == member_names(source_path)
    except Exception:
        media_equal, critical_equal, members_equal, media_count = False, False, False, 0
    non_table_equal = non_table_signatures(candidate) == non_table_signatures(source)
    notes_equal = notes_signature(candidate) == notes_signature(source)
    layouts_equal = (
        len(candidate.slides) == len(source.slides) == 1
        and str(candidate.slides[0].slide_layout.part.partname) == str(source.slides[0].slide_layout.part.partname)
    )
    preservation_passed = media_equal and media_count == 4 and critical_equal and non_table_equal and notes_equal and layouts_equal
    criteria.append(result_item(
        "R006",
        preservation_passed,
        f"非表格对象保持={non_table_equal}，备注={notes_equal}，版式={layouts_equal}，媒体={media_equal}/{media_count}，关键部件={critical_equal}，成员集={members_equal}",
    ))

    table_visual_readable = header_passed and banding_passed
    criteria.append(result_item(
        "R007",
        render_passed and table_visual_readable,
        f"表外相似度={similarity:.6f}，证据={json.dumps(similarity_evidence, ensure_ascii=False)}，{render_note}",
    ))

    shape_counts = [len(slide.shapes) for slide in candidate.slides]
    table_count = sum(1 for slide in candidate.slides for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.TABLE)
    editable_passed = (
        package_passed
        and len(candidate.slides) == 1
        and shape_counts == [3]
        and table_count == 1
        and media_equal
        and media_count == 4
        and members_equal
    )
    criteria.append(result_item("R008", editable_passed, f"对象数={shape_counts}，原生表格数={table_count}，媒体={media_count}，成员集={members_equal}"))

    artifact_files = sorted(str(item.relative_to(artifact_dir)) for item in artifact_dir.rglob("*") if item.is_file())
    all_text = "\n".join(
        shape.text
        for slide in candidate.slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )
    if candidate_inventory.get("passed"):
        all_text += "\n" + "\n".join(value for row in candidate_inventory["text"] for value in row)
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker.lower() in all_text.lower()]
    hygiene_passed = artifact_files == [OUTPUT_NAME] and not markers
    criteria.append(result_item("R009", hygiene_passed, f"输出文件={artifact_files}，可疑标记={markers}"))

    try:
        normalized_equal = normalized_slide_hash(path) == normalized_slide_hash(source_path)
        scaffold_equal = scaffold_hashes(path) == scaffold_hashes(source_path)
    except Exception:
        normalized_equal = scaffold_equal = False
    criteria.append(result_item("R010", normalized_equal and critical_equal and scaffold_equal, f"规范化 slide1={normalized_equal}，关键部件={critical_equal}，包骨架={scaffold_equal}"))

    severe = (
        not package_passed
        or not table_passed
        or not header_passed
        or not banding_passed
        or not theme_passed
        or not preservation_passed
        or not render_passed
        or not editable_passed
        or bool(markers)
        or sha256(path) == INPUT_SHA256
    )
    criteria.append(result_item("R011", not severe, f"严重错误={severe}，未修改原件={sha256(path) == INPUT_SHA256}"))
    order = {identifier: index for index, identifier in enumerate(SPECS)}
    criteria.sort(key=lambda item: order[item["criterion_id"]])
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    output = args.artifact_dir / OUTPUT_NAME
    if output.is_file():
        evaluation = evaluate(output, args.artifact_dir)
    else:
        evaluation = {
            "task_id": TASK_ID,
            "rubric_version": RUBRIC_VERSION,
            "judge_version": JUDGE_VERSION,
            "artifact_sha256": [],
            "criterion_results": [],
            "reward": 0.0,
            "passed": False,
            "hurdle_failures": ["missing_output"],
        }
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
