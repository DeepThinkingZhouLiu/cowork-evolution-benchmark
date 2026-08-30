#!/usr/bin/env python3
"""PPTArena Case 27 的独立合法解、边界例和负例构造。"""

from __future__ import annotations

import argparse
import hashlib
import io
import shutil
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Inches


OUTPUT_NAME = "FixImagePorportions_GroundTruthA.pptx"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"p": P_NS, "a": A_NS, "r": R_NS}
EMU_PER_POINT = 12_700


def qn(namespace: str, local: str) -> str:
    return f"{{{namespace}}}{local}"


def target(outdir: Path, case: str) -> Path:
    path = outdir / case / OUTPUT_NAME
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def serialize(root) -> bytes:
    return etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)


def rewrite_archive(source: Path, output: Path, transform) -> None:
    with zipfile.ZipFile(source) as archive_in, zipfile.ZipFile(output, "w") as archive_out:
        for info in archive_in.infolist():
            archive_out.writestr(info, transform(info.filename, archive_in.read(info.filename)))


def picture(root, shape_id: str):
    nodes = root.xpath(f"./p:cSld/p:spTree/p:pic[p:nvPicPr/p:cNvPr/@id='{shape_id}']", namespaces=NS)
    if len(nodes) != 1:
        raise ValueError(f"图片 {shape_id} 不唯一")
    return nodes[0]


def shape(root, shape_id: str):
    nodes = root.xpath(f"./p:cSld/p:spTree/p:sp[p:nvSpPr/p:cNvPr/@id='{shape_id}']", namespaces=NS)
    if len(nodes) != 1:
        raise ValueError(f"形状 {shape_id} 不唯一")
    return nodes[0]


def set_geometry(node, *, left: float | None = None, top: float | None = None, width: float | None = None, height: float | None = None, rotation: float | None = None) -> None:
    transform = node.find("./p:spPr/a:xfrm", NS)
    offset = transform.find("a:off", NS)
    extent = transform.find("a:ext", NS)
    if left is not None:
        offset.set("x", str(round(left * EMU_PER_POINT)))
    if top is not None:
        offset.set("y", str(round(top * EMU_PER_POINT)))
    if width is not None:
        extent.set("cx", str(round(width * EMU_PER_POINT)))
    if height is not None:
        extent.set("cy", str(round(height * EMU_PER_POINT)))
    if rotation is not None:
        if abs(rotation) < 1e-9:
            transform.attrib.pop("rot", None)
        else:
            transform.set("rot", str(round(rotation * 60_000)))


def layout_variant(source: Path, output: Path, positions: dict[str, tuple[float, float, float, float]]) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name != "ppt/slides/slide1.xml":
            return data
        root = etree.fromstring(data)
        for shape_id, values in positions.items():
            left, top, width, height = values
            set_geometry(picture(root, shape_id), left=left, top=top, width=width, height=height, rotation=0.0)
        return serialize(root)

    rewrite_archive(source, output, transform)


def mutate_slide(source: Path, output: Path, mutate) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name != "ppt/slides/slide1.xml":
            return data
        root = etree.fromstring(data)
        mutate(root)
        return serialize(root)

    rewrite_archive(source, output, transform)


def alter_picture(source: Path, output: Path, shape_id: str, **geometry) -> None:
    mutate_slide(source, output, lambda root: set_geometry(picture(root, shape_id), **geometry))


def remove_picture(source: Path, output: Path, shape_id: str) -> None:
    def mutate(root) -> None:
        node = picture(root, shape_id)
        node.getparent().remove(node)

    mutate_slide(source, output, mutate)


def swap_positions(source: Path, output: Path, first_id: str, second_id: str) -> None:
    def mutate(root) -> None:
        first = picture(root, first_id).find("./p:spPr/a:xfrm/a:off", NS)
        second = picture(root, second_id).find("./p:spPr/a:xfrm/a:off", NS)
        first_x, second_x = first.get("x"), second.get("x")
        first.set("x", second_x)
        second.set("x", first_x)

    mutate_slide(source, output, mutate)


def swap_contents(source: Path, output: Path, first_id: str, second_id: str) -> None:
    def mutate(root) -> None:
        first = picture(root, first_id).find("./p:blipFill/a:blip", NS)
        second = picture(root, second_id).find("./p:blipFill/a:blip", NS)
        key = qn(R_NS, "embed")
        first_id_value, second_id_value = first.get(key), second.get(key)
        first.set(key, second_id_value)
        second.set(key, first_id_value)

    mutate_slide(source, output, mutate)


def change_crop(source: Path, output: Path) -> None:
    def mutate(root) -> None:
        image = picture(root, "4")
        crop = image.find("./p:blipFill/a:srcRect", NS)
        crop.set("l", "22000")
        crop.set("r", "3000")

    mutate_slide(source, output, mutate)


def change_text(source: Path, output: Path, old: str, new: str) -> None:
    def mutate(root) -> None:
        nodes = [node for node in root.xpath(".//a:t", namespaces=NS) if node.text == old]
        if len(nodes) != 1:
            raise ValueError(f"文字 {old!r} 不唯一")
        nodes[0].text = new

    mutate_slide(source, output, mutate)


def move_label(source: Path, output: Path) -> None:
    presentation = Presentation(source)
    label = next(item for item in presentation.slides[0].shapes if int(item.shape_id) == 23)
    label.left = round(382.46 * EMU_PER_POINT)
    presentation.save(output)


def repack_content_types(source: Path, output: Path) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name != "[Content_Types].xml":
            return data
        root = etree.fromstring(data)
        root.append(etree.Comment("harmless repack"))
        return serialize(root)

    rewrite_archive(source, output, transform)


def change_theme(source: Path, output: Path) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name != "ppt/theme/theme1.xml":
            return data
        root = etree.fromstring(data)
        node = root.find("./a:themeElements/a:clrScheme/a:accent1/a:srgbClr", NS)
        node.set("val", "FF00FF")
        return serialize(root)

    rewrite_archive(source, output, transform)


def change_layout(source: Path, output: Path) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name != "ppt/slideLayouts/slideLayout12.xml":
            return data
        root = etree.fromstring(data)
        root.append(etree.Comment("unexpected layout mutation"))
        return serialize(root)

    rewrite_archive(source, output, transform)


def change_media(source: Path, output: Path) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name != "ppt/media/image1.png":
            return data
        image = Image.open(io.BytesIO(data)).convert("RGBA")
        red, green, blue, alpha = image.getpixel((0, 0))
        image.putpixel((0, 0), ((red + 1) % 256, green, blue, alpha))
        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        return buffer.getvalue()

    rewrite_archive(source, output, transform)


def add_textbox(source: Path, output: Path, text: str) -> None:
    presentation = Presentation(source)
    box = presentation.slides[0].shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(2.4), Inches(0.35))
    box.text = text
    presentation.save(output)


def add_slide(source: Path, output: Path) -> None:
    presentation = Presentation(source)
    presentation.slides.add_slide(presentation.slide_layouts[0])
    presentation.save(output)


def flatten_images(source: Path, output: Path) -> None:
    presentation = Presentation(source)
    slide = presentation.slides[0]
    image_shapes = []
    for item in slide.shapes:
        try:
            hashlib.sha256(item.image.blob).hexdigest()
        except Exception:
            continue
        image_shapes.append(item)
    for item in image_shapes:
        item.element.getparent().remove(item.element)
    with tempfile.TemporaryDirectory(prefix="pptarena027-flat-") as temporary:
        image_path = Path(temporary) / "flat.png"
        Image.new("RGB", (1600, 900), (240, 244, 250)).save(image_path)
        slide.shapes.add_picture(str(image_path), 0, 0, presentation.slide_width, presentation.slide_height)
        presentation.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--upstream-ground-truth", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()

    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    target(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, target(args.outdir, "unmodified"))
    shutil.copy2(args.upstream_ground_truth, target(args.outdir, "upstream-ground-truth-layout-mutated"))

    layout_variant(args.source, target(args.outdir, "valid-224-centered"), {"10": (89.34, 174.0, 224.0, 224.0), "4": (370.40, 174.0, 224.0, 224.0), "7": (651.46, 174.0, 224.0, 224.0)})
    layout_variant(args.source, target(args.outdir, "valid-232-centered"), {"10": (85.34, 165.0, 232.0, 232.0), "4": (366.40, 165.0, 232.0, 232.0), "7": (647.46, 165.0, 232.0, 232.0)})
    repack_content_types(args.golden, target(args.outdir, "repacked-content-types"))

    remove_picture(args.golden, target(args.outdir, "missing-image"), "7")
    swap_positions(args.golden, target(args.outdir, "wrong-order"), "10", "7")
    swap_contents(args.golden, target(args.outdir, "swapped-image-content"), "10", "7")
    alter_picture(args.golden, target(args.outdir, "unequal-size"), "7", width=215.0, height=215.0)
    alter_picture(args.golden, target(args.outdir, "not-square"), "4", width=235.0, height=220.0)
    alter_picture(args.golden, target(args.outdir, "misaligned-top"), "4", top=183.0)
    alter_picture(args.golden, target(args.outdir, "rotated-shark"), "10", rotation=5.0)
    alter_picture(args.golden, target(args.outdir, "overlap-label"), "7", top=180.0, width=240.0, height=240.0)
    alter_picture(args.golden, target(args.outdir, "off-center-car"), "4", left=380.0)
    alter_picture(args.golden, target(args.outdir, "uneven-gaps"), "4", left=381.0)
    change_crop(args.golden, target(args.outdir, "changed-crop"))
    change_text(args.golden, target(args.outdir, "changed-label"), "shark", "sharks")
    move_label(args.golden, target(args.outdir, "moved-label"))
    change_text(args.golden, target(args.outdir, "changed-title"), "What I like", "Things I like")
    change_media(args.golden, target(args.outdir, "changed-media"))
    change_theme(args.golden, target(args.outdir, "changed-theme"))
    change_layout(args.golden, target(args.outdir, "changed-layout"))
    flatten_images(args.golden, target(args.outdir, "flattened-images"))
    add_textbox(args.golden, target(args.outdir, "extra-object"), "extra object")
    add_textbox(args.golden, target(args.outdir, "hardcoded-marker"), "solution/Golden")
    add_slide(args.golden, target(args.outdir, "extra-slide"))


if __name__ == "__main__":
    main()
