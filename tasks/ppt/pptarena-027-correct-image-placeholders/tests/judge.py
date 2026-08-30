#!/usr/bin/env python3
"""PPTArena Case 27 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import math
import re
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from lxml import etree
from PIL import Image, ImageChops, ImageDraw
from pptx import Presentation


OUTPUT_NAME = "FixImagePorportions_GroundTruthA.pptx"
INPUT_NAME = "FixImagePorportions_TestA.pptx"
INPUT_SHA256 = "026285dad008ed551128b04548909d8e4377324b67f8428cb6a68d2c87d9555f"
TASK_ID = "cowork-evo/pptarena-027-correct-image-placeholders"
RUBRIC_VERSION = "pptarena-case-027-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-027-deterministic-v0.1"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
EMU_PER_POINT = 12_700
TARGETS = {
    "b47eaefbb5744156c6a8ec9e07c9ebf2e4c50fd6b1ac1eed660b5b76aae56931": "shark",
    "7a0a10e4a6c0bd69ca0bf6cdfb5262e2e6fde089ca5afc65144d5664ac1d722d": "car",
    "a5d6368aef46ca944c6b600782e62e8d382f12a1d8c78c16eb0f98fb42cd0e19": "baseball",
}
SUSPICIOUS_MARKERS = ["solution/golden", "/root/", "judge-result", "artifact_dir", "pptarena-case-027"]
SPECS = {
    "R001": ("指定 PPTX 可打开、页面尺寸不变且只有 1 页", 6, "hurdle"),
    "R002": ("鲨鱼、汽车和棒球原图与三个标签正确配对", 20, "hurdle"),
    "R003": ("三图无旋转、等大近方形、顶边对齐且间距均衡", 18, "hurdle"),
    "R004": ("三图中心与对应标签对齐且留有清晰空隙", 14, "hurdle"),
    "R005": ("标题、标签、样式、顺序和备注保持", 12, "hurdle"),
    "R006": ("主题、母版、版式、背景、超链接、媒体和裁剪保持", 12, "hurdle"),
    "R007": ("完整渲染 1 页且三图区外画面保持", 8, "hurdle"),
    "R008": ("保留 7 个原生对象、3 个图片占位符和 3 个媒体", 5, "hurdle"),
    "R009": ("输出目录干净且无答案标记", 3, "positive"),
    "R010": ("规范化三图合法几何后内容与包骨架保持", 2, "positive"),
    "R011": ("严重图片配对、布局、可读性、可编辑性或保持错误", -20, "penalty"),
}


def qn(namespace: str, local: str) -> str:
    return f"{{{namespace}}}{local}"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def locate_source() -> Path | None:
    candidates = [Path("/app/input_files") / INPUT_NAME, Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME]
    return next((path for path in candidates if path.is_file()), None)


def package_ok(path: Path) -> tuple[bool, str]:
    try:
        with zipfile.ZipFile(path) as archive:
            broken = archive.testzip()
            names = {name for name in archive.namelist() if not name.endswith("/")}
            required = {"[Content_Types].xml", "ppt/presentation.xml", "ppt/slides/slide1.xml", "ppt/theme/theme1.xml"}
            missing = sorted(required - names)
            macros = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
        return broken is None and not missing and not macros, f"损坏成员={broken}，缺失={missing}，宏={macros}"
    except Exception as exc:
        return False, f"ZIP 检查失败：{exc}"


def image_hash(shape) -> str | None:
    try:
        return hashlib.sha256(shape.image.blob).hexdigest()
    except Exception:
        return None


def image_crop(shape) -> tuple[float, float, float, float]:
    return tuple(round(float(value), 6) for value in (shape.crop_left, shape.crop_top, shape.crop_right, shape.crop_bottom))


def target_inventory(presentation: Presentation) -> tuple[dict[str, dict], list[str]]:
    values: dict[str, dict] = {}
    all_images = []
    if len(presentation.slides) != 1:
        return values, all_images
    for z_index, shape in enumerate(presentation.slides[0].shapes):
        digest = image_hash(shape)
        if digest is None:
            continue
        all_images.append(digest)
        if digest not in TARGETS:
            continue
        if digest in values:
            values[digest]["duplicate"] = True
            continue
        rotation = float(shape.rotation or 0) % 360
        rotation = min(rotation, abs(360 - rotation))
        values[digest] = {
            "semantic": TARGETS[digest],
            "shape_id": int(shape.shape_id),
            "name": shape.name,
            "z_index": z_index,
            "left_pt": shape.left / EMU_PER_POINT,
            "top_pt": shape.top / EMU_PER_POINT,
            "width_pt": shape.width / EMU_PER_POINT,
            "height_pt": shape.height / EMU_PER_POINT,
            "right_pt": (shape.left + shape.width) / EMU_PER_POINT,
            "bottom_pt": (shape.top + shape.height) / EMU_PER_POINT,
            "center_x_pt": (shape.left + shape.width / 2) / EMU_PER_POINT,
            "rotation": rotation,
            "crop": image_crop(shape),
            "duplicate": False,
        }
    return values, all_images


def label_inventory(presentation: Presentation) -> dict[str, dict]:
    values = {}
    if len(presentation.slides) != 1:
        return values
    for shape in presentation.slides[0].shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        text = shape.text.strip().lower()
        if text in {"shark", "car", "baseball"}:
            values[text] = {
                "shape_id": int(shape.shape_id),
                "left_pt": shape.left / EMU_PER_POINT,
                "top_pt": shape.top / EMU_PER_POINT,
                "width_pt": shape.width / EMU_PER_POINT,
                "height_pt": shape.height / EMU_PER_POINT,
                "center_x_pt": (shape.left + shape.width / 2) / EMU_PER_POINT,
            }
    return values


def geometry_assessment(targets: dict[str, dict], labels: dict[str, dict]) -> tuple[bool, bool, dict]:
    if set(targets) != set(TARGETS) or set(labels) != {"shark", "car", "baseball"}:
        return False, False, {"reason": "图片或标签不完整"}
    ordered = sorted(targets.values(), key=lambda item: item["left_pt"])
    semantics = [item["semantic"] for item in ordered]
    widths = [item["width_pt"] for item in ordered]
    heights = [item["height_pt"] for item in ordered]
    tops = [item["top_pt"] for item in ordered]
    rotations = [item["rotation"] for item in ordered]
    horizontal_gaps = [ordered[index + 1]["left_pt"] - ordered[index]["right_pt"] for index in range(2)]
    mapping_passed = semantics == ["shark", "car", "baseball"] and not any(item["duplicate"] for item in ordered)
    layout_passed = (
        all(215.0 <= value <= 240.0 for value in widths + heights)
        and max(widths) - min(widths) <= 2.0
        and max(heights) - min(heights) <= 2.0
        and max(abs(width - height) for width, height in zip(widths, heights)) <= 1.5
        and max(tops) - min(tops) <= 2.0
        and 155.0 <= min(tops) <= 185.0
        and max(rotations) <= 0.1
        and all(30.0 <= gap <= 75.0 for gap in horizontal_gaps)
        and max(horizontal_gaps) - min(horizontal_gaps) <= 6.0
        and ordered[0]["left_pt"] >= 45.0
        and ordered[-1]["right_pt"] <= 915.0
    )
    center_offsets = []
    vertical_gaps = []
    for item in ordered:
        label = labels[item["semantic"]]
        center_offsets.append(abs(item["center_x_pt"] - label["center_x_pt"]))
        vertical_gaps.append(label["top_pt"] - item["bottom_pt"])
    pairing_passed = (
        max(center_offsets) <= 4.0
        and all(5.0 <= gap <= 30.0 for gap in vertical_gaps)
        and max(vertical_gaps) - min(vertical_gaps) <= 5.0
    )
    return mapping_passed, layout_passed and pairing_passed, {
        "semantics": semantics,
        "widths_pt": widths,
        "heights_pt": heights,
        "tops_pt": tops,
        "rotations": rotations,
        "horizontal_gaps_pt": horizontal_gaps,
        "center_offsets_pt": center_offsets,
        "vertical_gaps_pt": vertical_gaps,
        "layout_passed": layout_passed,
        "pairing_passed": pairing_passed,
    }


def shape_hash(shape) -> str:
    return hashlib.sha256(etree.tostring(shape.element, method="c14n")).hexdigest()


def non_target_signatures(presentation: Presentation) -> list[tuple]:
    if len(presentation.slides) != 1:
        return []
    values = []
    for index, shape in enumerate(presentation.slides[0].shapes):
        if image_hash(shape) in TARGETS:
            continue
        values.append((index, int(shape.shape_id), shape.name, str(shape.shape_type), shape_hash(shape)))
    return values


def notes_signature(presentation: Presentation) -> tuple[str, ...]:
    values = []
    for slide in presentation.slides:
        try:
            values.append(slide.notes_slide.notes_text_frame.text)
        except Exception:
            values.append("")
    return tuple(values)


def media_hashes(path: Path) -> Counter[str]:
    with zipfile.ZipFile(path) as archive:
        return Counter(hashlib.sha256(archive.read(name)).hexdigest() for name in archive.namelist() if name.startswith("ppt/media/") and not name.endswith("/"))


def member_names(path: Path) -> set[str]:
    with zipfile.ZipFile(path) as archive:
        return {name for name in archive.namelist() if not name.endswith("/")}


def critical_hashes(path: Path) -> dict[str, str]:
    prefixes = ("ppt/theme/", "ppt/slideMasters/", "ppt/slideLayouts/", "ppt/notesMasters/", "ppt/notesSlides/", "ppt/media/", "ppt/slides/_rels/")
    with zipfile.ZipFile(path) as archive:
        return {name: hashlib.sha256(archive.read(name)).hexdigest() for name in sorted(archive.namelist()) if not name.endswith("/") and name.startswith(prefixes)}


def scaffold_hashes(path: Path) -> dict[str, str]:
    parts = {"[Content_Types].xml", "_rels/.rels", "ppt/presentation.xml", "ppt/_rels/presentation.xml.rels"}
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        return {name: hashlib.sha256(archive.read(name)).hexdigest() for name in sorted(parts & names)}


def normalized_slide_hash(path: Path, source: Path) -> tuple[str, str]:
    with zipfile.ZipFile(path) as candidate_archive, zipfile.ZipFile(source) as source_archive:
        candidate_root = etree.fromstring(candidate_archive.read("ppt/slides/slide1.xml"))
        source_root = etree.fromstring(source_archive.read("ppt/slides/slide1.xml"))
    for shape_id in ("10", "4", "7"):
        source_nodes = source_root.xpath(f"./p:cSld/p:spTree/p:pic[p:nvPicPr/p:cNvPr/@id='{shape_id}']/p:spPr/a:xfrm", namespaces={"p": P_NS, "a": A_NS})
        candidate_nodes = candidate_root.xpath(f"./p:cSld/p:spTree/p:pic[p:nvPicPr/p:cNvPr/@id='{shape_id}']/p:spPr/a:xfrm", namespaces={"p": P_NS, "a": A_NS})
        if len(source_nodes) != 1 or len(candidate_nodes) != 1:
            raise ValueError(f"无法规范化图片 {shape_id}")
        candidate_nodes[0].getparent().replace(candidate_nodes[0], copy.deepcopy(source_nodes[0]))
    source_common = source_root.find(qn(P_NS, "cSld"))
    candidate_common = candidate_root.find(qn(P_NS, "cSld"))
    return (
        hashlib.sha256(etree.tostring(source_common, method="c14n")).hexdigest(),
        hashlib.sha256(etree.tostring(candidate_common, method="c14n")).hexdigest(),
    )


def render_presentation(path: Path, outdir: Path) -> tuple[list[Path], int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "lo-profile"
    profile.mkdir()
    completed = subprocess.run(["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(path)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=240, check=False)
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return [], 0, f"转 PDF 失败：{completed.stdout[-500:]}"
    info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=30, check=False).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    pages = int(match.group(1)) if match else 0
    raster = subprocess.run(["pdftoppm", "-jpeg", "-r", "150", str(pdf), str(outdir / "slide")], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=240, check=False)
    images = sorted(outdir.glob("slide-*.jpg")) if raster.returncode == 0 else []
    return images, pages, f"页数={pages}，渲染图={len(images)}，soffice={completed.returncode}，pdftoppm={raster.returncode}"


def masked_similarity(source_image: Image.Image, candidate_image: Image.Image, source_targets: dict[str, dict], candidate_targets: dict[str, dict], slide_size: tuple[int, int]) -> tuple[float, dict]:
    source = source_image.convert("RGB")
    candidate = candidate_image.convert("RGB")
    if source.size != candidate.size:
        return 0.0, {"reason": "渲染尺寸不一致"}
    mask = Image.new("L", source.size, 255)
    draw = ImageDraw.Draw(mask)
    boxes = []
    padding = max(10, round(min(source.size) * 0.008))
    for inventory in (source_targets, candidate_targets):
        for item in inventory.values():
            left = round(item["left_pt"] * source.width / (slide_size[0] / EMU_PER_POINT))
            top = round(item["top_pt"] * source.height / (slide_size[1] / EMU_PER_POINT))
            right = round(item["right_pt"] * source.width / (slide_size[0] / EMU_PER_POINT))
            bottom = round(item["bottom_pt"] * source.height / (slide_size[1] / EMU_PER_POINT))
            box = (max(0, left - padding), max(0, top - padding), min(source.width, right + padding), min(source.height, bottom + padding))
            boxes.append(box)
            draw.rectangle(box, fill=0)
    difference = ImageChops.difference(source, candidate)
    pixels = mask.histogram()[255]
    histogram = difference.histogram(mask)
    squared = sum(count * ((index % 256) ** 2) for index, count in enumerate(histogram))
    rms = math.sqrt(squared / (pixels * 3)) if pixels else 255.0
    return max(0.0, 1.0 - rms / 255.0), {"masked_boxes": boxes, "evaluated_pixels": pixels, "rms": round(rms, 6)}


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {"criterion_id": identifier, "score": 1.0 if passed else 0.0, "evidence": evidence, "method": "deterministic", "raw": {"description": description, "weight": weight, "type": kind, "passed": bool(passed)}}


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [item["criterion_id"] for item in criteria if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0 and item["raw"]["passed"])
    penalty = sum(abs(item["raw"]["weight"]) for item in criteria if item["raw"]["weight"] < 0 and not item["raw"]["passed"])
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [sha256(path)] if path.is_file() else [], "criterion_results": criteria, "reward": reward, "passed": reward >= 0.8 and not hurdles, "hurdle_failures": hurdles}


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        return make_result(path, [result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配") for identifier in SPECS])
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        return make_result(path, [result_item(identifier, False, f"PPTX 打开失败：{exc}") for identifier in SPECS])

    criteria = []
    package_passed, package_evidence = package_ok(path)
    size_passed = len(candidate.slides) == len(source.slides) == 1 and candidate.slide_width == source.slide_width and candidate.slide_height == source.slide_height
    criteria.append(result_item("R001", package_passed and size_passed, f"{package_evidence}，页数={len(candidate.slides)}，页面={candidate.slide_width}×{candidate.slide_height}"))

    source_targets, source_images = target_inventory(source)
    candidate_targets, candidate_images = target_inventory(candidate)
    labels = label_inventory(candidate)
    mapping_passed, combined_geometry, geometry = geometry_assessment(candidate_targets, labels)
    media_set_matches = Counter(candidate_images) == Counter(TARGETS.keys())
    criteria.append(result_item("R002", mapping_passed and media_set_matches, f"图片数={len(candidate_images)}，语义={geometry.get('semantics')}，媒体哈希集匹配={media_set_matches}"))
    layout_passed = bool(geometry.get("layout_passed"))
    pairing_passed = bool(geometry.get("pairing_passed"))
    criteria.append(result_item("R003", layout_passed, json.dumps(geometry, ensure_ascii=False)))
    criteria.append(result_item("R004", pairing_passed, json.dumps({"center_offsets_pt": geometry.get("center_offsets_pt"), "vertical_gaps_pt": geometry.get("vertical_gaps_pt")}, ensure_ascii=False)))

    non_target_equal = non_target_signatures(candidate) == non_target_signatures(source)
    notes_equal = notes_signature(candidate) == notes_signature(source)
    criteria.append(result_item("R005", non_target_equal and notes_equal, f"非目标对象签名保持={non_target_equal}，备注={notes_equal}"))

    try:
        media_equal = media_hashes(path) == media_hashes(source_path)
        media_count = sum(media_hashes(path).values())
        critical_equal = critical_hashes(path) == critical_hashes(source_path)
        members_equal = member_names(path) == member_names(source_path)
    except Exception:
        media_equal, critical_equal, members_equal, media_count = False, False, False, 0
    crops_equal = set(candidate_targets) == set(source_targets) and all(candidate_targets[key]["crop"] == source_targets[key]["crop"] for key in source_targets)
    criteria.append(result_item("R006", media_equal and media_count == 3 and critical_equal and crops_equal and members_equal, f"媒体={media_equal}/{media_count}，裁剪={crops_equal}，主题母版版式与关系={critical_equal}，成员集={members_equal}"))

    similarity = 0.0
    similarity_evidence = {}
    render_note = "前置检查失败，未执行渲染"
    render_passed = False
    if package_passed and size_passed and len(candidate_targets) == 3:
        with tempfile.TemporaryDirectory(prefix="pptarena027-judge-") as temporary:
            root = Path(temporary)
            source_renders, source_pages, source_note = render_presentation(source_path, root / "source")
            candidate_renders, candidate_pages, candidate_note = render_presentation(path, root / "candidate")
            render_note = f"源[{source_note}]；候选[{candidate_note}]"
            if len(source_renders) == len(candidate_renders) == source_pages == candidate_pages == 1:
                with Image.open(source_renders[0]) as source_image, Image.open(candidate_renders[0]) as candidate_image:
                    similarity, similarity_evidence = masked_similarity(source_image, candidate_image, source_targets, candidate_targets, (source.slide_width, source.slide_height))
                render_passed = similarity >= 0.9980
    criteria.append(result_item("R007", render_passed and layout_passed and pairing_passed, f"图区外相似度={similarity:.6f}，{json.dumps(similarity_evidence, ensure_ascii=False)}，{render_note}"))

    shape_counts = [len(slide.shapes) for slide in candidate.slides]
    image_count = sum(1 for slide in candidate.slides for shape in slide.shapes if image_hash(shape) is not None)
    editable_passed = package_passed and shape_counts == [7] and image_count == 3 and len(candidate_targets) == 3 and media_equal and media_count == 3 and members_equal
    criteria.append(result_item("R008", editable_passed, f"对象数={shape_counts}，已填充图片占位符={image_count}，媒体={media_count}，成员集={members_equal}"))

    artifact_files = sorted(str(item.relative_to(artifact_dir)) for item in artifact_dir.rglob("*") if item.is_file())
    all_text = "\n".join(shape.text for slide in candidate.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False))
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker.lower() in all_text.lower()]
    hygiene_passed = artifact_files == [OUTPUT_NAME] and not markers
    criteria.append(result_item("R009", hygiene_passed, f"输出文件={artifact_files}，可疑标记={markers}"))

    try:
        source_normalized, candidate_normalized = normalized_slide_hash(path, source_path)
        normalized_equal = source_normalized == candidate_normalized
        scaffold_equal = scaffold_hashes(path) == scaffold_hashes(source_path)
    except Exception:
        normalized_equal = scaffold_equal = False
    criteria.append(result_item("R010", normalized_equal and critical_equal and scaffold_equal, f"规范化 slide1={normalized_equal}，关键部件={critical_equal}，包骨架={scaffold_equal}"))

    severe = not all([package_passed, mapping_passed, layout_passed, pairing_passed, non_target_equal, notes_equal, media_equal, critical_equal, crops_equal, render_passed, editable_passed]) or bool(markers) or sha256(path) == INPUT_SHA256
    criteria.append(result_item("R011", not severe, f"严重错误={severe}，未修改原件={sha256(path) == INPUT_SHA256}"))
    order = {identifier: index for index, identifier in enumerate(SPECS)}
    criteria.sort(key=lambda item: order[item["criterion_id"]])
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    output = args.artifact_dir / OUTPUT_NAME
    if output.is_file():
        evaluation = evaluate(output, args.artifact_dir)
    else:
        evaluation = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [], "criterion_results": [], "reward": 0.0, "passed": False, "hurdle_failures": ["missing_output"]}
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
