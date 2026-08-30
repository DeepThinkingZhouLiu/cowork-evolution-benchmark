#!/usr/bin/env python3
"""PPTArena Case 28 的独立合法变体、边界例和负例。"""

from __future__ import annotations

import argparse
import hashlib
import io
import shutil
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Inches


OUTPUT_NAME = "FixAgendaAndOrder_GroundTruthA.pptx"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A16_NS = "http://schemas.microsoft.com/office/drawing/2014/main"
NS = {"p": P_NS, "a": A_NS, "r": R_NS, "a16": A16_NS}


def target(outdir: Path, case: str) -> Path:
    path = outdir / case / OUTPUT_NAME
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def serialize(root) -> bytes:
    return etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)


def rewrite(source: Path, output: Path, transform) -> None:
    with zipfile.ZipFile(source) as archive_in, zipfile.ZipFile(output, "w") as archive_out:
        for info in archive_in.infolist():
            archive_out.writestr(info, transform(info.filename, archive_in.read(info.filename)))


def mutate_xml(source: Path, output: Path, part: str, mutate) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name != part:
            return data
        root = etree.fromstring(data)
        mutate(root)
        return serialize(root)
    rewrite(source, output, transform)


def valid_identifier_variant(source: Path, output: Path) -> None:
    def mutate(root) -> None:
        shapes = root.xpath("./p:cSld/p:spTree/p:sp[p:nvSpPr/p:nvPr/p:ph/@type='sldNum']", namespaces=NS)
        node = shapes[0]
        properties = node.find("./p:nvSpPr/p:cNvPr", NS)
        properties.set("id", "28")
        properties.set("name", "Slide Number Placeholder 28")
        creation = properties.find("./a:extLst/a:ext/a16:creationId", NS)
        if creation is not None:
            creation.set("id", "{28000000-0000-0000-0000-000000000003}")
        field = node.find("./p:txBody/a:p/a:fld", NS)
        field.set("id", "{28000000-0000-0000-0000-000000000013}")
    mutate_xml(source, output, "ppt/slides/slide4.xml", mutate)


def valid_whitespace_variant(source: Path, output: Path) -> None:
    def mutate(root) -> None:
        for node in root.xpath("./p:cSld/p:spTree/p:sp[p:nvSpPr/p:cNvPr/@id='4']//a:t", namespaces=NS):
            if node.text:
                node.text = f" {node.text.strip()} "
    mutate_xml(source, output, "ppt/slides/slide2.xml", mutate)


def repack_content_types(source: Path, output: Path) -> None:
    mutate_xml(source, output, "[Content_Types].xml", lambda root: root.append(etree.Comment("harmless repack")))


def change_order(source: Path, output: Path) -> None:
    def mutate(root) -> None:
        slide_list = root.find("p:sldIdLst", NS)
        nodes = list(slide_list)
        slide_list.remove(nodes[0])
        slide_list.insert(1, nodes[0])
    mutate_xml(source, output, "ppt/presentation.xml", mutate)


def agenda_variant(source: Path, output: Path, replacements: list[str]) -> None:
    def mutate(root) -> None:
        text_nodes = root.xpath("./p:cSld/p:spTree/p:sp[p:nvSpPr/p:cNvPr/@id='4']/p:txBody/a:p//a:t", namespaces=NS)
        for node, value in zip(text_nodes, replacements):
            node.text = value
    mutate_xml(source, output, "ppt/slides/slide2.xml", mutate)


def set_number(source: Path, output: Path, part: int, value: str) -> None:
    def mutate(root) -> None:
        nodes = root.xpath("./p:cSld/p:spTree/p:sp[p:nvSpPr/p:nvPr/p:ph/@type='sldNum']//a:fld[@type='slidenum']/a:t", namespaces=NS)
        nodes[0].text = value
    mutate_xml(source, output, f"ppt/slides/slide{part}.xml", mutate)


def remove_number(source: Path, output: Path, part: int) -> None:
    def mutate(root) -> None:
        nodes = root.xpath("./p:cSld/p:spTree/p:sp[p:nvSpPr/p:nvPr/p:ph/@type='sldNum']", namespaces=NS)
        nodes[0].getparent().remove(nodes[0])
    mutate_xml(source, output, f"ppt/slides/slide{part}.xml", mutate)


def change_text(source: Path, output: Path, part: int, old: str, new: str) -> None:
    def mutate(root) -> None:
        nodes = [node for node in root.xpath(".//a:t", namespaces=NS) if node.text and old in node.text]
        if len(nodes) != 1:
            raise ValueError(f"文本 {old!r} 不唯一：{len(nodes)}")
        nodes[0].text = nodes[0].text.replace(old, new)
    mutate_xml(source, output, f"ppt/slides/slide{part}.xml", mutate)


def change_theme(source: Path, output: Path) -> None:
    def mutate(root) -> None:
        node = root.xpath(".//*[local-name()='accent1']/*[1]")[0]
        node.set("val", "FF00FF")
    mutate_xml(source, output, "ppt/theme/theme1.xml", mutate)


def change_layout(source: Path, output: Path) -> None:
    mutate_xml(source, output, "ppt/slideLayouts/slideLayout34.xml", lambda root: root.append(etree.Comment("unexpected layout mutation")))


def change_relationship(source: Path, output: Path) -> None:
    mutate_xml(source, output, "ppt/slides/_rels/slide4.xml.rels", lambda root: root.append(etree.Comment("unexpected relationship mutation")))


def change_media(source: Path, output: Path) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name != "ppt/media/image1.png":
            return data
        image = Image.open(io.BytesIO(data)).convert("RGBA")
        pixel = image.getpixel((0, 0))
        image.putpixel((0, 0), ((pixel[0] + 1) % 256, pixel[1], pixel[2], pixel[3]))
        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        return buffer.getvalue()
    rewrite(source, output, transform)


def add_textbox(source: Path, output: Path, text: str) -> None:
    presentation = Presentation(source)
    box = presentation.slides[0].shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(2.2), Inches(0.3))
    box.text = text
    presentation.save(output)


def add_slide(source: Path, output: Path) -> None:
    presentation = Presentation(source)
    presentation.slides.add_slide(presentation.slide_layouts[0])
    presentation.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--upstream-ground-truth", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()

    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    target(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, target(args.outdir, "unmodified"))
    shutil.copy2(args.upstream_ground_truth, target(args.outdir, "upstream-ground-truth-package-mutated"))
    valid_identifier_variant(args.golden, target(args.outdir, "valid-number-identifiers"))
    valid_whitespace_variant(args.golden, target(args.outdir, "valid-agenda-whitespace"))
    repack_content_types(args.golden, target(args.outdir, "repacked-content-types"))
    set_number(args.golden, target(args.outdir, "wrong-number"), 2, "9")
    remove_number(args.golden, target(args.outdir, "missing-number-3"), 4)
    change_order(args.golden, target(args.outdir, "wrong-order"))
    agenda_variant(args.golden, target(args.outdir, "wrong-agenda-order"), ["Conclusion", "Q&A", "Visual appeal", "Engaging the audience", "Effective delivery techniques"])
    agenda_variant(args.golden, target(args.outdir, "duplicate-agenda-item"), ["Visual appeal", "Engaging the audience", "Effective delivery techniques", "Conclusion", "Conclusion"])
    change_text(args.golden, target(args.outdir, "changed-title"), 1, "Why visuals matter", "Why visual design matters")
    change_text(args.golden, target(args.outdir, "changed-body"), 4, "Make eye contact", "Maintain eye contact")
    change_text(args.golden, target(args.outdir, "changed-date"), 6, "8/13/25", "8/14/25")
    change_theme(args.golden, target(args.outdir, "changed-theme"))
    change_layout(args.golden, target(args.outdir, "changed-layout"))
    change_relationship(args.golden, target(args.outdir, "changed-relationship"))
    change_media(args.golden, target(args.outdir, "changed-media"))
    add_textbox(args.golden, target(args.outdir, "extra-object"), "extra object")
    add_textbox(args.golden, target(args.outdir, "hardcoded-marker"), "solution/Golden")
    add_slide(args.golden, target(args.outdir, "extra-slide"))


if __name__ == "__main__":
    main()
