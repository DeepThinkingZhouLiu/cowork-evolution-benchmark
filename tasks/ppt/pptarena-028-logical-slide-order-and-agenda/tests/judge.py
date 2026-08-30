#!/usr/bin/env python3
"""PPTArena Case 28 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from PIL import Image, ImageChops, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "FixAgendaAndOrder_GroundTruthA.pptx"
INPUT_NAME = "FixAgendaAndOrder_TestA.pptx"
INPUT_SHA256 = "bcbdf13b305029f1fca7c33d0b6b0020f84acf963bf066471d4be00971c2175b"
TASK_ID = "cowork-evo/pptarena-028-logical-slide-order-and-agenda"
RUBRIC_VERSION = "pptarena-case-028-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-028-deterministic-v0.1"
EMU_PER_POINT = 12_700
TITLE_ORDER = ["Agenda", "Why visuals matter", "Engaging the audience", "Effective delivery techniques", "Conclusion", "Q&A"]
SOURCE_INDEX_ORDER = [1, 0, 3, 5, 2, 4]
AGENDA_ORDER = ["Visual appeal", "Engaging the audience", "Effective delivery techniques", "Conclusion", "Q&A"]
RELATIONSHIP_ORDER = ["rId3", "rId2", "rId5", "rId7", "rId4", "rId6"]
EXPECTED_SHAPE_COUNTS = [5, 5, 5, 5, 9, 5]
SUSPICIOUS_MARKERS = ["solution/golden", "/root/", "judge-result", "artifact_dir", "pptarena-case-028"]
SPECS = {
    "R001": ("可打开、页尺寸不变且严格 6 页", 6, "hurdle"),
    "R002": ("六页按指定逻辑顺序排列", 20, "hurdle"),
    "R003": ("议程五项完整且顺序与后续页一致", 16, "hurdle"),
    "R004": ("六页可编辑页码为 1–6 且位置样式一致", 18, "positive"),
    "R005": ("非目标内容、样式、几何和备注保持", 14, "hurdle"),
    "R006": ("媒体、关系、主题、母版、版式和包成员保持", 10, "hurdle"),
    "R007": ("六页可渲染且非目标区域视觉保持", 8, "hurdle"),
    "R008": ("原生对象数正确且无压平或额外页", 4, "hurdle"),
    "R009": ("输出目录干净且无答案标记", 2, "positive"),
    "R010": ("原关系 ID 与包骨架保持", 2, "positive"),
    "R011": ("无严重页序、议程、页码、保持或可编辑性错误", -20, "penalty"),
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalize(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def locate_source() -> Path | None:
    candidates = [Path("/app/input_files") / INPUT_NAME, Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME]
    return next((path for path in candidates if path.is_file()), None)


def package_ok(path: Path) -> tuple[bool, str]:
    try:
        with zipfile.ZipFile(path) as archive:
            broken = archive.testzip()
            names = {name for name in archive.namelist() if not name.endswith("/")}
            required = {"[Content_Types].xml", "ppt/presentation.xml"} | {f"ppt/slides/slide{i}.xml" for i in range(1, 7)}
            missing = sorted(required - names)
            macros = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
        return broken is None and not missing and not macros, f"损坏成员={broken}，缺失={missing}，宏={macros}"
    except Exception as exc:
        return False, f"ZIP 检查失败：{exc}"


def placeholder_type(shape) -> str:
    if not getattr(shape, "is_placeholder", False):
        return ""
    try:
        return str(shape.placeholder_format.type)
    except Exception:
        return ""


def is_slide_number(shape) -> bool:
    return "SLIDE_NUMBER" in placeholder_type(shape)


def slide_title(slide) -> str:
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        text = normalize(shape.text)
        if text in TITLE_ORDER:
            return text
    return ""


def agenda_items(slide) -> list[str]:
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False) or normalize(shape.text) == "Agenda":
            continue
        items = [normalize(paragraph.text) for paragraph in shape.text_frame.paragraphs if normalize(paragraph.text)]
        if set(items).intersection(AGENDA_ORDER):
            return items
    return []


def page_number(shape) -> str:
    return normalize(shape.text) if getattr(shape, "has_text_frame", False) else ""


def page_number_assessment(presentation: Presentation) -> tuple[bool, dict]:
    values = []
    geometries = []
    for index, slide in enumerate(presentation.slides, 1):
        shapes = [shape for shape in slide.shapes if is_slide_number(shape)]
        values.append([page_number(shape) for shape in shapes])
        if len(shapes) == 1:
            shape = shapes[0]
            rotation = float(shape.rotation or 0) % 360
            geometries.append({
                "left_pt": round(shape.left / EMU_PER_POINT, 2),
                "top_pt": round(shape.top / EMU_PER_POINT, 2),
                "width_pt": round(shape.width / EMU_PER_POINT, 2),
                "height_pt": round(shape.height / EMU_PER_POINT, 2),
                "rotation": round(rotation, 2),
            })
        else:
            geometries.append({})
    exact = values == [[str(index)] for index in range(1, 7)]
    geometry_ok = len(geometries) == 6 and all(
        910 <= item.get("left_pt", -1) <= 930
        and 490 <= item.get("top_pt", -1) <= 510
        and 45 <= item.get("width_pt", -1) <= 60
        and 25 <= item.get("height_pt", -1) <= 40
        and min(abs(item.get("rotation", 0)), abs(item.get("rotation", 0) - 270), abs(item.get("rotation", 0) + 90)) <= 1
        for item in geometries
    )
    return exact and geometry_ok, {"values": values, "geometries": geometries, "exact": exact, "geometry_ok": geometry_ok}


def font_signature(font) -> tuple:
    try:
        rgb = str(font.color.rgb) if font.color.rgb is not None else None
    except Exception:
        rgb = None
    try:
        theme = int(font.color.theme_color) if font.color.theme_color is not None else None
    except Exception:
        theme = None
    return (font.name, round(font.size.pt, 3) if font.size else None, font.bold, font.italic, font.underline, rgb, theme)


def shape_signature(shape, *, agenda: bool = False) -> tuple:
    text = ""
    paragraphs = ()
    if getattr(shape, "has_text_frame", False):
        text = "<AGENDA>" if agenda else normalize(shape.text)
        paragraph_values = []
        for paragraph in shape.text_frame.paragraphs:
            paragraph_values.append((
                "<ITEM>" if agenda and normalize(paragraph.text) else normalize(paragraph.text),
                paragraph.level,
                str(paragraph.alignment),
                str(paragraph.line_spacing),
                tuple(font_signature(run.font) for run in paragraph.runs if normalize(run.text)),
            ))
        paragraphs = tuple(paragraph_values)
    picture = ""
    crop = ()
    try:
        picture = hashlib.sha256(shape.image.blob).hexdigest()
        crop = tuple(round(float(value), 6) for value in (shape.crop_left, shape.crop_top, shape.crop_right, shape.crop_bottom))
    except Exception:
        pass
    children = ()
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        children = tuple(shape_signature(child) for child in shape.shapes)
    return (
        str(shape.shape_type), placeholder_type(shape), shape.name,
        tuple(round(int(value) / 1000) * 1000 for value in (shape.left, shape.top, shape.width, shape.height)),
        round(float(shape.rotation or 0), 3), text, paragraphs, picture, crop, children,
    )


def slide_signature(slide, title: str) -> tuple:
    values = []
    for shape in slide.shapes:
        if is_slide_number(shape):
            continue
        is_agenda = title == "Agenda" and getattr(shape, "has_text_frame", False) and set(
            normalize(p.text) for p in shape.text_frame.paragraphs if normalize(p.text)
        ) == set(AGENDA_ORDER)
        values.append(shape_signature(shape, agenda=is_agenda))
    return tuple(values)


def notes_signature(slide) -> tuple[str, ...]:
    try:
        return tuple(normalize(shape.text) for shape in slide.notes_slide.shapes if getattr(shape, "has_text_frame", False))
    except Exception:
        return ()


def media_hashes(path: Path) -> Counter[str]:
    with zipfile.ZipFile(path) as archive:
        return Counter(hashlib.sha256(archive.read(name)).hexdigest() for name in archive.namelist() if name.startswith("ppt/media/") and not name.endswith("/"))


def member_names(path: Path) -> set[str]:
    with zipfile.ZipFile(path) as archive:
        return {name for name in archive.namelist() if not name.endswith("/")}


def critical_hashes(path: Path) -> dict[str, str]:
    prefixes = ("ppt/theme/", "ppt/slideMasters/", "ppt/slideLayouts/", "ppt/notesMasters/", "ppt/notesSlides/", "ppt/media/", "ppt/slides/_rels/")
    with zipfile.ZipFile(path) as archive:
        return {name: hashlib.sha256(archive.read(name)).hexdigest() for name in sorted(archive.namelist()) if not name.endswith("/") and name.startswith(prefixes)}


def presentation_relationship_order(path: Path) -> list[str]:
    from lxml import etree
    namespace = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main", "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    with zipfile.ZipFile(path) as archive:
        root = etree.fromstring(archive.read("ppt/presentation.xml"))
    return root.xpath("./p:sldIdLst/p:sldId/@r:id", namespaces=namespace)


def scaffold_hashes(path: Path) -> dict[str, str]:
    names = {"[Content_Types].xml", "_rels/.rels", "ppt/_rels/presentation.xml.rels"}
    with zipfile.ZipFile(path) as archive:
        return {name: hashlib.sha256(archive.read(name)).hexdigest() for name in sorted(names)}


def render(path: Path, outdir: Path) -> tuple[list[Path], str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "lo-profile"
    profile.mkdir()
    completed = subprocess.run(["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(path)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=240, check=False)
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return [], f"转 PDF 失败：{completed.stdout[-400:]}"
    raster = subprocess.run(["pdftoppm", "-jpeg", "-r", "100", str(pdf), str(outdir / "slide")], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=240, check=False)
    images = sorted(outdir.glob("slide-*.jpg")) if raster.returncode == 0 else []
    return images, f"soffice={completed.returncode}，pdftoppm={raster.returncode}，渲染页={len(images)}"


def masked_similarity(source_path: Path, candidate_path: Path, agenda: bool) -> float:
    with Image.open(source_path) as source_image, Image.open(candidate_path) as candidate_image:
        source = source_image.convert("RGB")
        candidate = candidate_image.convert("RGB")
        if source.size != candidate.size:
            return 0.0
        mask = Image.new("L", source.size, 255)
        draw = ImageDraw.Draw(mask)
        width, height = source.size
        draw.rectangle((int(width * 0.93), int(height * 0.88), width, height), fill=0)
        if agenda:
            draw.rectangle((int(width * 0.52), int(height * 0.40), int(width * 0.95), int(height * 0.96)), fill=0)
        difference = ImageChops.difference(source, candidate)
        pixels = mask.histogram()[255]
        histogram = difference.histogram(mask)
        squared = sum(count * ((index % 256) ** 2) for index, count in enumerate(histogram))
        rms = math.sqrt(squared / (pixels * 3)) if pixels else 255.0
        return max(0.0, 1.0 - rms / 255.0)


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {"criterion_id": identifier, "score": 1.0 if passed else 0.0, "evidence": evidence, "method": "deterministic", "raw": {"description": description, "weight": weight, "type": kind, "passed": bool(passed)}}


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [item["criterion_id"] for item in criteria if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0 and item["raw"]["passed"])
    penalty = sum(abs(item["raw"]["weight"]) for item in criteria if item["raw"]["weight"] < 0 and not item["raw"]["passed"])
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [sha256(path)] if path.is_file() else [], "criterion_results": criteria, "reward": reward, "passed": reward >= 0.8 and not hurdles, "hurdle_failures": hurdles}


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        return make_result(path, [result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配") for identifier in SPECS])
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        return make_result(path, [result_item(identifier, False, f"PPTX 打开失败：{exc}") for identifier in SPECS])

    criteria = []
    package_passed, package_note = package_ok(path)
    size_passed = len(source.slides) == len(candidate.slides) == 6 and source.slide_width == candidate.slide_width and source.slide_height == candidate.slide_height
    criteria.append(result_item("R001", package_passed and size_passed, f"{package_note}，页数={len(candidate.slides)}，页面={candidate.slide_width}×{candidate.slide_height}"))

    titles = [slide_title(slide) for slide in candidate.slides]
    order_passed = titles == TITLE_ORDER
    criteria.append(result_item("R002", order_passed, f"实际顺序={titles}"))

    agenda = agenda_items(candidate.slides[0]) if len(candidate.slides) else []
    agenda_passed = agenda == AGENDA_ORDER
    criteria.append(result_item("R003", agenda_passed, f"实际议程={agenda}"))

    numbers_passed, number_evidence = page_number_assessment(candidate)
    criteria.append(result_item("R004", numbers_passed, json.dumps(number_evidence, ensure_ascii=False)))

    signature_mismatches = []
    notes_mismatches = []
    if len(candidate.slides) == 6:
        for logical_index, source_index in enumerate(SOURCE_INDEX_ORDER):
            title = TITLE_ORDER[logical_index]
            if slide_signature(candidate.slides[logical_index], title) != slide_signature(source.slides[source_index], title):
                signature_mismatches.append(logical_index + 1)
            if notes_signature(candidate.slides[logical_index]) != notes_signature(source.slides[source_index]):
                notes_mismatches.append(logical_index + 1)
    else:
        signature_mismatches = list(range(1, 7))
        notes_mismatches = list(range(1, 7))
    content_preserved = not signature_mismatches and not notes_mismatches
    criteria.append(result_item("R005", content_preserved, f"内容/样式变化页={signature_mismatches}，备注变化页={notes_mismatches}"))

    try:
        media_equal = media_hashes(path) == media_hashes(source_path)
        critical_equal = critical_hashes(path) == critical_hashes(source_path)
        members_equal = member_names(path) == member_names(source_path)
    except Exception:
        media_equal = critical_equal = members_equal = False
    package_preserved = media_equal and critical_equal and members_equal
    criteria.append(result_item("R006", package_preserved, f"媒体={media_equal}，主题/母版/版式/关系={critical_equal}，成员集={members_equal}"))

    similarities = []
    render_note = "前置检查失败，未执行渲染"
    if package_passed and size_passed and order_passed and content_preserved and package_preserved:
        with tempfile.TemporaryDirectory(prefix="pptarena028-judge-") as temporary:
            root = Path(temporary)
            source_images, source_note = render(source_path, root / "source")
            candidate_images, candidate_note = render(path, root / "candidate")
            render_note = f"源[{source_note}]；候选[{candidate_note}]"
            if len(source_images) == len(candidate_images) == 6:
                for logical_index, source_index in enumerate(SOURCE_INDEX_ORDER):
                    similarities.append(masked_similarity(source_images[source_index], candidate_images[logical_index], logical_index == 0))
    render_passed = len(similarities) == 6 and min(similarities) >= 0.998
    criteria.append(result_item("R007", render_passed, f"遮罩后逐页相似度={[round(value, 6) for value in similarities]}，{render_note}"))

    shape_counts = [len(slide.shapes) for slide in candidate.slides]
    full_slide_pictures = []
    for index, slide in enumerate(candidate.slides, 1):
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.width * shape.height >= 0.85 * candidate.slide_width * candidate.slide_height:
                full_slide_pictures.append((index, shape.name))
    editable_passed = shape_counts == EXPECTED_SHAPE_COUNTS and not full_slide_pictures
    criteria.append(result_item("R008", editable_passed, f"对象数={shape_counts}，全页压平图={full_slide_pictures}"))

    artifact_files = sorted(str(item.relative_to(artifact_dir)) for item in artifact_dir.rglob("*") if item.is_file())
    all_text = "\n".join(shape.text for slide in candidate.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False))
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker.lower() in all_text.lower()]
    hygiene_passed = artifact_files == [OUTPUT_NAME] and not markers
    criteria.append(result_item("R009", hygiene_passed, f"输出文件={artifact_files}，可疑标记={markers}"))

    try:
        relationship_order = presentation_relationship_order(path)
        scaffold_equal = scaffold_hashes(path) == scaffold_hashes(source_path)
    except Exception:
        relationship_order, scaffold_equal = [], False
    skeleton_passed = relationship_order == RELATIONSHIP_ORDER and scaffold_equal
    criteria.append(result_item("R010", skeleton_passed, f"sldId 关系顺序={relationship_order}，包骨架={scaffold_equal}"))

    severe = (
        not all([package_passed, size_passed, order_passed, agenda_passed, content_preserved, package_preserved, render_passed, editable_passed])
        or bool(markers)
        or not numbers_passed
        or sha256(path) == INPUT_SHA256
    )
    criteria.append(result_item("R011", not severe, f"严重错误={severe}，未修改原件={sha256(path) == INPUT_SHA256}"))
    order = {identifier: index for index, identifier in enumerate(SPECS)}
    criteria.sort(key=lambda item: order[item["criterion_id"]])
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    output = args.artifact_dir / OUTPUT_NAME
    if output.is_file():
        evaluation = evaluate(output, args.artifact_dir)
    else:
        evaluation = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [], "criterion_results": [], "reward": 0.0, "passed": False, "hurdle_failures": ["missing_output"]}
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
