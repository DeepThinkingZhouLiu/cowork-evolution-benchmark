#!/usr/bin/env python3
"""只为 PPTArena Case 29 构造可重放的合法变体、负例和中等例。"""

from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt


OUTPUT_NAME = "PieChart_GroundTruthA.pptx"
TARGET_SLIDE_INDEX = 2


def output_path(root: Path, case: str) -> Path:
    target = root / case / OUTPUT_NAME
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def exact_chart(presentation: Presentation, slide_index: int = TARGET_SLIDE_INDEX):
    charts = [
        shape
        for shape in presentation.slides[slide_index].shapes
        if shape.shape_type == MSO_SHAPE_TYPE.CHART
    ]
    if len(charts) != 1:
        raise ValueError(f"第 {slide_index + 1} 页应有唯一图表，实际为 {len(charts)}")
    return charts[0]


def chart_data_from(shape) -> ChartData:
    source_chart = shape.chart
    categories = list(source_chart.plots[0].categories)
    series = list(source_chart.series)
    if len(categories) != 8 or len(series) != 1:
        raise ValueError("冻结输入应包含 8 个类别和 1 个数据系列")
    data = ChartData()
    data.categories = categories
    data.add_series(series[0].name, list(series[0].values))
    return data


def replace_target_chart(
    presentation: Presentation,
    chart_type,
    left,
    top,
    width,
    height,
):
    slide = presentation.slides[TARGET_SLIDE_INDEX]
    old_chart = exact_chart(presentation)
    data = chart_data_from(old_chart)
    old_chart._element.getparent().remove(old_chart._element)
    return slide.shapes.add_chart(
        chart_type,
        left,
        top,
        width,
        height,
        data,
    ).chart


def valid_alternate(source: Path, output: Path) -> None:
    """从冻结输入重建不依赖 Golden 坐标与标签样式的等价原生饼图。"""
    presentation = Presentation(source)
    chart = replace_target_chart(
        presentation,
        XL_CHART_TYPE.PIE,
        Inches(0.80),
        Inches(1.90),
        Inches(8.40),
        Inches(4.75),
    )
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.RIGHT
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(10)
    plot = chart.plots[0]
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.show_category_name = False
    labels.show_value = False
    labels.show_percentage = True
    labels.font.size = Pt(10)
    presentation.save(output)


def wrong_chart_type(source: Path, output: Path) -> None:
    presentation = Presentation(source)
    chart = replace_target_chart(
        presentation,
        XL_CHART_TYPE.LINE_MARKERS,
        Inches(0.80),
        Inches(1.90),
        Inches(8.40),
        Inches(4.75),
    )
    chart.has_legend = False
    presentation.save(output)


def wrong_value(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    chart = exact_chart(presentation).chart
    values = chart._element.xpath(
        "./c:chart/c:plotArea/c:pieChart/c:ser/c:val//c:pt[@idx='0']/c:v"
    )
    if len(values) != 1:
        raise ValueError("无法定位饼图第一个数值")
    values[0].text = "0.9"
    presentation.save(output)


def wrong_category(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    chart = exact_chart(presentation).chart
    categories = chart._element.xpath(
        "./c:chart/c:plotArea/c:pieChart/c:ser/c:cat//c:pt[@idx='7']/c:v"
    )
    if len(categories) != 1:
        raise ValueError("无法定位饼图第八个类别")
    categories[0].text = "Other"
    presentation.save(output)


def second_chart(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    data = ChartData()
    data.categories = ["A", "B"]
    data.add_series("extra", [1, 2])
    presentation.slides[TARGET_SLIDE_INDEX].shapes.add_chart(
        XL_CHART_TYPE.PIE,
        Inches(7.70),
        Inches(5.00),
        Inches(1.50),
        Inches(1.20),
        data,
    )
    presentation.save(output)


def changed_slide3_title(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    slide = presentation.slides[TARGET_SLIDE_INDEX]
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        if "Asian and NHOPI populations" in shape.text:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if "Asian and NHOPI populations" in run.text:
                        run.text = run.text.replace("diverse groups", "one group")
                        presentation.save(output)
                        return
    raise ValueError("无法定位第 3 页主标题")


def changed_other_chart(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    chart = exact_chart(presentation, slide_index=4).chart
    values = chart._element.xpath(
        ".//c:val//c:numCache/c:pt[@idx='0']/c:v | "
        ".//c:yVal//c:numCache/c:pt[@idx='0']/c:v"
    )
    if not values:
        raise ValueError("无法定位第 5 页图表数值")
    values[0].text = str(float(values[0].text) + 0.25)
    presentation.save(output)


def flattened_slide3(golden: Path, slide_image: Path, output: Path) -> None:
    presentation = Presentation(golden)
    slide = presentation.slides[TARGET_SLIDE_INDEX]
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    slide.shapes.add_picture(
        str(slide_image),
        0,
        0,
        width=presentation.slide_width,
        height=presentation.slide_height,
    )
    presentation.save(output)


def hardcoded_marker(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    box = presentation.slides[TARGET_SLIDE_INDEX].shapes.add_textbox(
        Inches(0.10), Inches(7.10), Inches(2.50), Inches(0.20)
    )
    box.text = "solution/Golden"
    presentation.save(output)


def no_labels_or_legend(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    chart = exact_chart(presentation).chart
    chart.has_legend = False
    for node in chart._element.xpath(".//c:dLbls"):
        for tag in ["showCatName", "showVal", "showPercent"]:
            children = node.xpath(f"./c:{tag}")
            if children:
                children[0].set("val", "0")
    presentation.save(output)


def tiny_labels(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    chart = exact_chart(presentation).chart
    sized = chart._element.xpath(".//c:dLbls//*[@sz]")
    if not sized:
        raise ValueError("无法定位数据标签字号")
    for node in sized:
        node.set("sz", "400")
    presentation.save(output)


def undersized_chart(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    chart = exact_chart(presentation)
    chart.left = Inches(3.25)
    chart.top = Inches(2.30)
    chart.width = Inches(3.50)
    chart.height = Inches(2.50)
    presentation.save(output)


def extra_slide(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    # 该上游文稿只保留了一个自定义版式，不假定存在常规第 7 个空白版式。
    presentation.slides.add_slide(presentation.slide_layouts[0])
    presentation.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--slide3-image", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()

    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))
    valid_alternate(args.source, output_path(args.outdir, "valid-alternate"))
    wrong_chart_type(args.source, output_path(args.outdir, "wrong-chart-type"))
    wrong_value(args.golden, output_path(args.outdir, "wrong-value"))
    wrong_category(args.golden, output_path(args.outdir, "wrong-category"))
    second_chart(args.golden, output_path(args.outdir, "second-chart"))
    changed_slide3_title(args.golden, output_path(args.outdir, "changed-slide3-title"))
    changed_other_chart(args.golden, output_path(args.outdir, "changed-other-chart"))
    flattened_slide3(
        args.golden,
        args.slide3_image,
        output_path(args.outdir, "flattened-slide3"),
    )
    hardcoded_marker(args.golden, output_path(args.outdir, "hardcoded-marker"))
    no_labels_or_legend(args.golden, output_path(args.outdir, "no-labels-or-legend"))
    tiny_labels(args.golden, output_path(args.outdir, "tiny-labels"))
    undersized_chart(args.golden, output_path(args.outdir, "undersized-chart"))
    extra_slide(args.golden, output_path(args.outdir, "extra-slide"))


if __name__ == "__main__":
    main()
