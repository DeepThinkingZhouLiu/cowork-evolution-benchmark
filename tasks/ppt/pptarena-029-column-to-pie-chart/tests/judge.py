#!/usr/bin/env python3
"""PPTArena Case 29 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import statistics
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from PIL import Image, ImageChops, ImageDraw, ImageStat
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn


OUTPUT_NAME = "PieChart_GroundTruthA.pptx"
INPUT_NAME = "PieChart_TestA.pptx"
INPUT_SHA256 = "07911f768a13823f9e8a4a426a7d2667a16f8283a4237f2883abe8b24527415e"
TASK_ID = "cowork-evo/pptarena-029-column-to-pie-chart"
RUBRIC_VERSION = "pptarena-case-029-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-029-deterministic-v0.1"
EMU_PER_INCH = 914_400
PAGE_EDGE_TOLERANCE = int(0.03 * EMU_PER_INCH)
TARGET_SLIDE_INDEX = 2
ALLOWED_PIE_TYPES = {XL_CHART_TYPE.PIE, XL_CHART_TYPE.PIE_EXPLODED}
SUSPICIOUS_MARKERS = [
    "solution/golden",
    "/root/",
    "judge-result",
    "artifact_dir",
    "pptarena-case-029",
]

SPECS = {
    "R001": ("指定 PPTX 可打开，页面尺寸不变且总页数为 9", 6, "hurdle"),
    "R002": ("第 3 页恰有一个原生可编辑饼图", 14, "hurdle"),
    "R003": ("饼图单系列、8 个类别和底层数值与原图一致", 18, "hurdle"),
    "R004": ("除目标图表外的页面、内容、图表、图片、顺序和备注保持", 14, "hurdle"),
    "R005": ("第 3 页的全部非图表对象、内容、样式和几何保持", 10, "hurdle"),
    "R006": ("饼图位于主内容区、尺寸足够且不与其他对象重叠", 8, "positive"),
    "R007": ("全部 8 个类别可辨识，并显示数值或百分比", 8, "hurdle"),
    "R008": ("非目标页与第 3 页非图表区的渲染效果保持", 8, "hurdle"),
    "R009": ("图表和其他页面可编辑，无压平、宏、损坏关系或媒体丢失", 7, "hurdle"),
    "R010": ("9 页可渲染，第 3 页饼图可见、未裁切且非空图", 5, "positive"),
    "R011": ("输出目录干净且不含答案路径或未授权文字", 2, "positive"),
    "R012": ("页数、图表类型、数据、原页或可编辑性无严重错误", -20, "penalty"),
}


def normalize_visible(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def locate_source() -> Path | None:
    candidates = [
        Path("/app/input_files") / INPUT_NAME,
        Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME,
    ]
    return next((path for path in candidates if path.is_file()), None)


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def placeholder_type(shape) -> str:
    if not getattr(shape, "is_placeholder", False):
        return ""
    return str(shape.placeholder_format.type)


def geometry(shape) -> tuple[int, int, int, int]:
    return (int(shape.left), int(shape.top), int(shape.width), int(shape.height))


def rounded_geometry(shape) -> tuple[int, int, int, int]:
    return tuple(round(value / 1000) * 1000 for value in geometry(shape))


def color_signature(font) -> tuple:
    color = font.color
    try:
        color_type = int(color.type) if color.type is not None else None
    except Exception:
        color_type = str(color.type)
    try:
        rgb = str(color.rgb) if color.rgb is not None else None
    except Exception:
        rgb = None
    try:
        theme = int(color.theme_color) if color.theme_color is not None else None
    except Exception:
        theme = None
    try:
        brightness = round(color.brightness, 5)
    except Exception:
        brightness = None
    return (color_type, rgb, theme, brightness)


def run_style_signature(font) -> tuple:
    return (
        font.name,
        round(font.size.pt, 4) if font.size else None,
        bool(font.bold),
        bool(font.italic),
        bool(font.underline),
        int(font.language_id) if font.language_id is not None else None,
        color_signature(font),
    )


def text_format_signature(shape) -> tuple:
    text_frame = shape.text_frame
    paragraphs = []
    for paragraph in text_frame.paragraphs:
        styles = {
            run_style_signature(run.font)
            for run in paragraph.runs
            if run.text.strip()
        }
        paragraphs.append(
            (
                normalize_visible(paragraph.text),
                paragraph.level,
                str(paragraph.alignment),
                str(paragraph.line_spacing),
                str(paragraph.space_before),
                str(paragraph.space_after),
                tuple(sorted(styles, key=repr)),
            )
        )
    return (
        int(shape.shape_type),
        placeholder_type(shape),
        rounded_geometry(shape),
        (
            text_frame.margin_left,
            text_frame.margin_right,
            text_frame.margin_top,
            text_frame.margin_bottom,
            str(text_frame.vertical_anchor),
            text_frame.word_wrap,
            str(text_frame.auto_size),
        ),
        tuple(paragraphs),
    )


def chart_categories(chart) -> tuple[str, ...]:
    # compare_chart_data() 传入的是图表 shape，其他调用点传入 Chart。
    # 这里同时兼容两种对象，避免把 OOXML 缓存中的真实类别误判为空。
    if hasattr(chart, "chart"):
        chart = chart.chart
    try:
        return tuple(normalize_visible(str(value)) for value in chart.plots[0].categories)
    except Exception:
        return ()


def numeric_signature(value):
    if value is None:
        return None
    try:
        return round(float(value), 12)
    except Exception:
        return str(value)


def chart_data_signature(shape) -> tuple:
    chart = shape.chart
    series = tuple(
        (
            normalize_visible(item.name or ""),
            tuple(numeric_signature(value) for value in item.values),
        )
        for item in chart.series
    )
    return (str(chart.chart_type), chart_categories(chart), series)


def shape_content_signature(shape) -> tuple:
    text = normalize_visible(shape.text) if getattr(shape, "has_text_frame", False) else ""
    picture_hash = ""
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        picture_hash = hashlib.sha256(shape.image.blob).hexdigest()
    chart = ()
    if shape.shape_type == MSO_SHAPE_TYPE.CHART:
        chart = chart_data_signature(shape)
    children = ()
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        children = tuple(shape_content_signature(child) for child in shape.shapes)
    return (
        str(shape.shape_type),
        placeholder_type(shape),
        text,
        picture_hash,
        chart,
        children,
    )


def shape_style_signature(shape) -> tuple:
    text_style = ()
    if getattr(shape, "has_text_frame", False):
        text_style = text_format_signature(shape)
    picture_hash = ""
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        picture_hash = hashlib.sha256(shape.image.blob).hexdigest()
    chart_style = ()
    if shape.shape_type == MSO_SHAPE_TYPE.CHART:
        chart_style = (
            chart_data_signature(shape),
            bool(shape.chart.has_legend),
            shape.chart.chart_style,
        )
    children = ()
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        children = tuple(shape_style_signature(child) for child in shape.shapes)
    return (
        str(shape.shape_type),
        placeholder_type(shape),
        rounded_geometry(shape),
        text_style,
        picture_hash,
        chart_style,
        children,
    )


def slide_content_signature(slide, include_charts: bool = True) -> tuple:
    return tuple(
        shape_content_signature(shape)
        for shape in slide.shapes
        if include_charts or shape.shape_type != MSO_SHAPE_TYPE.CHART
    )


def slide_style_signature(slide, include_charts: bool = True) -> tuple:
    return tuple(
        shape_style_signature(shape)
        for shape in slide.shapes
        if include_charts or shape.shape_type != MSO_SHAPE_TYPE.CHART
    )


def slide_notes(slide) -> tuple[str, ...]:
    try:
        return tuple(
            shape.text
            for shape in slide.notes_slide.shapes
            if getattr(shape, "has_text_frame", False)
        )
    except Exception:
        return ()


def media_hashes(path: Path) -> Counter[str]:
    with zipfile.ZipFile(path) as archive:
        return Counter(
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("ppt/media/")
        )


def package_ok(path: Path) -> tuple[bool, str]:
    try:
        with zipfile.ZipFile(path) as archive:
            broken = archive.testzip()
            names = set(archive.namelist())
            required = {
                "[Content_Types].xml",
                "ppt/presentation.xml",
                "ppt/slides/slide1.xml",
                "ppt/slides/slide9.xml",
            }
            missing = sorted(required - names)
            macros = sorted(
                name for name in names if name.lower().endswith("vbaproject.bin")
            )
        return broken is None and not missing and not macros, (
            f"损坏成员={broken}，缺失={missing}，宏={macros}"
        )
    except Exception as exc:
        return False, f"ZIP 检查失败：{exc}"


def chart_shapes(slide) -> list:
    return [
        shape
        for shape in iter_shapes(slide.shapes)
        if shape.shape_type == MSO_SHAPE_TYPE.CHART
    ]


def compare_chart_data(source_chart, candidate_chart) -> tuple[bool, dict]:
    expected_categories = chart_categories(source_chart)
    actual_categories = chart_categories(candidate_chart)
    expected_series = list(source_chart.chart.series)
    actual_series = list(candidate_chart.chart.series)
    evidence = {
        "expected_categories": list(expected_categories),
        "actual_categories": list(actual_categories),
        "expected_series_count": len(expected_series),
        "actual_series_count": len(actual_series),
        "category_value_mismatches": [],
        "series_name": "",
    }
    if len(expected_series) != 1 or len(actual_series) != 1:
        return False, evidence
    expected_values = list(expected_series[0].values)
    actual_values = list(actual_series[0].values)
    expected_name = normalize_visible(expected_series[0].name or "")
    actual_name = normalize_visible(actual_series[0].name or "")
    evidence["series_name"] = f"{expected_name!r}→{actual_name!r}"
    if (
        len(expected_categories) != 8
        or len(actual_categories) != 8
        or len(expected_values) != 8
        or len(actual_values) != 8
        or len(set(actual_categories)) != 8
        or set(expected_categories) != set(actual_categories)
    ):
        return False, evidence
    expected_map = dict(zip(expected_categories, expected_values, strict=True))
    actual_map = dict(zip(actual_categories, actual_values, strict=True))
    for category in expected_categories:
        expected = expected_map[category]
        actual = actual_map[category]
        equal = (
            expected is None
            and actual is None
            or expected is not None
            and actual is not None
            and math.isclose(
                float(expected),
                float(actual),
                rel_tol=1e-10,
                abs_tol=1e-12,
            )
        )
        if not equal:
            evidence["category_value_mismatches"].append(
                {"category": category, "expected": expected, "actual": actual}
            )
    passed = not evidence["category_value_mismatches"] and expected_name == actual_name
    return passed, evidence


def direct_chart_flag(parent, tag: str, default: bool = False) -> bool:
    element = parent.find(qn(f"c:{tag}")) if parent is not None else None
    if element is None:
        return default
    return str(element.get("val", "1")).lower() not in {"0", "false"}


def chart_label_evidence(chart_shape, category_count: int) -> tuple[bool, dict]:
    chart = chart_shape.chart
    chart_xml = chart._element
    plot_nodes = chart_xml.xpath(
        "./c:chart/c:plotArea/c:pieChart/c:dLbls | "
        "./c:chart/c:plotArea/c:pie3DChart/c:dLbls"
    )
    series_nodes = chart_xml.xpath(
        "./c:chart/c:plotArea/c:pieChart/c:ser/c:dLbls | "
        "./c:chart/c:plotArea/c:pie3DChart/c:ser/c:dLbls"
    )
    plot_labels = plot_nodes[0] if plot_nodes else None
    series_labels = series_nodes[0] if series_nodes else None
    # PowerPoint 可以把标签放在 plot 层，也可以放在 series 层。
    # Case 29 的人工 Golden 属于后者，python-pptx 新建图表则通常属于前者。
    global_category = direct_chart_flag(plot_labels, "showCatName")
    global_value = direct_chart_flag(plot_labels, "showVal")
    global_percent = direct_chart_flag(plot_labels, "showPercent")
    if series_labels is not None:
        global_category = direct_chart_flag(
            series_labels, "showCatName", global_category
        )
        global_value = direct_chart_flag(series_labels, "showVal", global_value)
        global_percent = direct_chart_flag(
            series_labels, "showPercent", global_percent
        )
    per_point: dict[int, tuple[bool, bool, bool]] = {}
    for dlabels in [plot_labels, series_labels]:
        if dlabels is None:
            continue
        for item in dlabels.xpath("./c:dLbl"):
            index_element = item.find(qn("c:idx"))
            if index_element is None:
                continue
            try:
                index = int(index_element.get("val"))
            except Exception:
                continue
            per_point[index] = (
                direct_chart_flag(item, "showCatName", global_category),
                direct_chart_flag(item, "showVal", global_value),
                direct_chart_flag(item, "showPercent", global_percent),
            )
    category_flags = [
        per_point.get(index, (global_category, global_value, global_percent))[0]
        for index in range(category_count)
    ]
    measure_flags = [
        any(per_point.get(index, (global_category, global_value, global_percent))[1:])
        for index in range(category_count)
    ]
    explicit_sizes = [
        int(value) / 100
        for value in chart_xml.xpath(
            "./c:chart/c:plotArea/c:pieChart/c:dLbls//@sz | "
            "./c:chart/c:plotArea/c:pieChart/c:ser/c:dLbls//@sz | "
            "./c:chart/c:plotArea/c:pie3DChart/c:dLbls//@sz | "
            "./c:chart/c:plotArea/c:pie3DChart/c:ser/c:dLbls//@sz | "
            "./c:chart/c:legend//@sz"
        )
        if str(value).isdigit()
    ]
    minimum_explicit_size = min(explicit_sizes) if explicit_sizes else None
    categories_identified = bool(chart.has_legend) or all(category_flags)
    measures_identified = all(measure_flags)
    readable_size = minimum_explicit_size is None or minimum_explicit_size >= 8.0
    passed = (
        category_count == 8
        and categories_identified
        and measures_identified
        and readable_size
    )
    return passed, {
        "has_legend": bool(chart.has_legend),
        "global_category": global_category,
        "global_value": global_value,
        "global_percent": global_percent,
        "category_flags": category_flags,
        "measure_flags": measure_flags,
        "minimum_explicit_font_pt": minimum_explicit_size,
    }


def chart_layout_checks(
    chart_shape,
    slide_width: int,
    slide_height: int,
) -> tuple[bool, dict]:
    if chart_shape is None:
        return False, {"reason": "缺失目标图表"}
    left, top, width, height = geometry(chart_shape)
    right = left + width
    bottom = top + height
    content_top = int(1.70 * EMU_PER_INCH)
    content_bottom = int(6.80 * EMU_PER_INCH)
    inside = (
        left >= int(0.05 * EMU_PER_INCH) - PAGE_EDGE_TOLERANCE
        and right <= slide_width - int(0.05 * EMU_PER_INCH) + PAGE_EDGE_TOLERANCE
        and top >= content_top - PAGE_EDGE_TOLERANCE
        and bottom <= content_bottom + PAGE_EDGE_TOLERANCE
    )
    large_enough = width >= 4.0 * EMU_PER_INCH and height >= 3.0 * EMU_PER_INCH
    return inside and large_enough, {
        "box_in": [
            round(left / EMU_PER_INCH, 4),
            round(top / EMU_PER_INCH, 4),
            round(width / EMU_PER_INCH, 4),
            round(height / EMU_PER_INCH, 4),
        ],
        "allowed_content_box_in": [0.05, 1.70, 9.90, 5.10],
        "inside": inside,
        "large_enough": large_enough,
    }


def full_slide_pictures(presentation: Presentation) -> list[str]:
    values: list[str] = []
    page_area = presentation.slide_width * presentation.slide_height
    for slide_index, slide in enumerate(presentation.slides, start=1):
        for shape in iter_shapes(slide.shapes):
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            coverage = shape.width * shape.height / page_area
            if coverage >= 0.85:
                values.append(f"slide{slide_index}:{shape.name}:{coverage:.3f}")
    return values


def render_presentation(path: Path, outdir: Path) -> tuple[list[Path], int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "lo-profile"
    profile.mkdir()
    completed = subprocess.run(
        [
            "soffice",
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(outdir),
            str(path),
        ],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return [], 0, f"转 PDF 失败：{completed.stdout[-500:]}"
    info = subprocess.run(
        ["pdfinfo", str(pdf)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=30,
        check=False,
    ).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    pages = int(match.group(1)) if match else 0
    prefix = outdir / "slide"
    raster = subprocess.run(
        ["pdftoppm", "-jpeg", "-r", "90", str(pdf), str(prefix)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    images = sorted(outdir.glob("slide-*.jpg")) if raster.returncode == 0 else []
    return images, pages, (
        f"页数={pages}，渲染图={len(images)}，"
        f"soffice={completed.returncode}，pdftoppm={raster.returncode}"
    )


def image_similarity(first: Path, second: Path) -> float:
    with Image.open(first) as left_image, Image.open(second) as right_image:
        left = left_image.convert("RGB")
        right = right_image.convert("RGB")
        if left.size != right.size:
            return 0.0
        difference = ImageChops.difference(left, right)
        histogram = difference.histogram()
        squared = sum(
            count * ((index % 256) ** 2)
            for index, count in enumerate(histogram)
        )
        pixels = left.width * left.height
        rms = math.sqrt(squared / (pixels * 3)) if pixels else 255.0
        return max(0.0, 1.0 - rms / 255.0)


def emu_rect_to_pixels(
    rectangle: tuple[int, int, int, int],
    image_size: tuple[int, int],
    slide_size: tuple[int, int],
    padding: int = 0,
) -> tuple[int, int, int, int]:
    image_width, image_height = image_size
    slide_width, slide_height = slide_size
    left, top, width, height = rectangle
    return (
        max(0, math.floor(left * image_width / slide_width) - padding),
        max(0, math.floor(top * image_height / slide_height) - padding),
        min(image_width, math.ceil((left + width) * image_width / slide_width) + padding),
        min(image_height, math.ceil((top + height) * image_height / slide_height) + padding),
    )


def masked_image_similarity(
    first: Path,
    second: Path,
    ignored_rectangles: list[tuple[int, int, int, int]],
    slide_size: tuple[int, int],
) -> float:
    with Image.open(first) as first_image, Image.open(second) as second_image:
        left = first_image.convert("RGB")
        right = second_image.convert("RGB")
        if left.size != right.size:
            return 0.0
        difference = ImageChops.difference(left, right)
        mask = Image.new("L", left.size, 255)
        draw = ImageDraw.Draw(mask)
        for rectangle in ignored_rectangles:
            draw.rectangle(
                emu_rect_to_pixels(rectangle, left.size, slide_size, padding=10),
                fill=0,
            )
        stats = ImageStat.Stat(difference, mask=mask)
        rms = statistics.mean(stats.rms) if stats.rms else 255.0
        return max(0.0, 1.0 - rms / 255.0)


def chart_crop_metrics(
    image_path: Path,
    chart_shape,
    slide_size: tuple[int, int],
) -> dict:
    if chart_shape is None:
        return {"stddev": 0.0, "nonwhite_ratio": 0.0}
    with Image.open(image_path) as image:
        rgb = image.convert("RGB")
        rectangle = emu_rect_to_pixels(geometry(chart_shape), rgb.size, slide_size)
        crop = rgb.crop(rectangle)
        if crop.width == 0 or crop.height == 0:
            return {"stddev": 0.0, "nonwhite_ratio": 0.0}
        stddev = statistics.mean(ImageStat.Stat(crop.convert("L")).stddev)
        histogram = crop.convert("L").histogram()
        nonwhite = sum(histogram[:245]) / (crop.width * crop.height)
        return {
            "stddev": float(stddev),
            "nonwhite_ratio": float(nonwhite),
            "crop_px": [crop.width, crop.height],
        }


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {
        "criterion_id": identifier,
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
        "raw": {
            "description": description,
            "weight": weight,
            "type": kind,
            "passed": bool(passed),
        },
    }


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [
        item["criterion_id"]
        for item in criteria
        if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]
    ]
    total = sum(
        item["raw"]["weight"]
        for item in criteria
        if item["raw"]["weight"] > 0
    )
    earned = sum(
        item["raw"]["weight"]
        for item in criteria
        if item["raw"]["weight"] > 0 and item["raw"]["passed"]
    )
    penalty = sum(
        abs(item["raw"]["weight"])
        for item in criteria
        if item["raw"]["weight"] < 0 and not item["raw"]["passed"]
    )
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [sha256(path)] if path.is_file() else [],
        "criterion_results": criteria,
        "reward": reward,
        "passed": reward >= 0.8 and not hurdles,
        "hurdle_failures": hurdles,
    }


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        criteria = [
            result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配")
            for identifier in SPECS
        ]
        return make_result(path, criteria)
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        criteria = [
            result_item(identifier, False, f"PPTX 打开失败：{exc}")
            for identifier in SPECS
        ]
        return make_result(path, criteria)

    criteria: list[dict] = []
    package_passed, package_evidence = package_ok(path)
    size_passed = (
        len(source.slides) == 9
        and len(candidate.slides) == 9
        and candidate.slide_width == source.slide_width
        and candidate.slide_height == source.slide_height
    )
    criteria.append(
        result_item(
            "R001",
            package_passed and size_passed,
            f"{package_evidence}，输入页数={len(source.slides)}，"
            f"候选页数={len(candidate.slides)}，页面="
            f"{candidate.slide_width / EMU_PER_INCH:.4f}×"
            f"{candidate.slide_height / EMU_PER_INCH:.4f}英寸",
        )
    )

    source_target_charts = chart_shapes(source.slides[TARGET_SLIDE_INDEX])
    candidate_target_charts = (
        chart_shapes(candidate.slides[TARGET_SLIDE_INDEX])
        if len(candidate.slides) > TARGET_SLIDE_INDEX
        else []
    )
    source_chart = source_target_charts[0] if len(source_target_charts) == 1 else None
    target_chart = candidate_target_charts[0] if len(candidate_target_charts) == 1 else None
    pie_passed = (
        len(candidate_target_charts) == 1
        and target_chart is not None
        and target_chart.chart.chart_type in ALLOWED_PIE_TYPES
    )
    criteria.append(
        result_item(
            "R002",
            pie_passed,
            f"第 3 页图表数={len(candidate_target_charts)}，类型="
            f"{str(target_chart.chart.chart_type) if target_chart is not None else '缺失'}，"
            f"对象类型={str(target_chart.shape_type) if target_chart is not None else '缺失'}",
        )
    )

    if source_chart is not None and target_chart is not None:
        data_passed, data_evidence = compare_chart_data(source_chart, target_chart)
    else:
        data_passed, data_evidence = False, {"reason": "源图表或候选图表缺失"}
    criteria.append(
        result_item(
            "R003",
            data_passed,
            f"图表数据={json.dumps(data_evidence, ensure_ascii=False)}",
        )
    )

    content_mismatches: list[int] = []
    style_mismatches: list[int] = []
    notes_mismatches: list[int] = []
    if len(candidate.slides) == 9:
        for index in range(9):
            if index == TARGET_SLIDE_INDEX:
                continue
            if slide_content_signature(candidate.slides[index]) != slide_content_signature(
                source.slides[index]
            ):
                content_mismatches.append(index + 1)
            if slide_style_signature(candidate.slides[index]) != slide_style_signature(
                source.slides[index]
            ):
                style_mismatches.append(index + 1)
            if slide_notes(candidate.slides[index]) != slide_notes(source.slides[index]):
                notes_mismatches.append(index + 1)
    else:
        content_mismatches = [1, 2, 4, 5, 6, 7, 8, 9]
        style_mismatches = content_mismatches[:]
        notes_mismatches = content_mismatches[:]
    criteria.append(
        result_item(
            "R004",
            not content_mismatches and not style_mismatches and not notes_mismatches,
            f"非目标页内容/图表变化={content_mismatches}，"
            f"样式/几何变化={style_mismatches}，备注变化={notes_mismatches}",
        )
    )

    target_nonchart_content_passed = False
    target_nonchart_style_passed = False
    target_notes_passed = False
    if len(candidate.slides) > TARGET_SLIDE_INDEX:
        target_nonchart_content_passed = slide_content_signature(
            candidate.slides[TARGET_SLIDE_INDEX], include_charts=False
        ) == slide_content_signature(
            source.slides[TARGET_SLIDE_INDEX], include_charts=False
        )
        target_nonchart_style_passed = slide_style_signature(
            candidate.slides[TARGET_SLIDE_INDEX], include_charts=False
        ) == slide_style_signature(
            source.slides[TARGET_SLIDE_INDEX], include_charts=False
        )
        target_notes_passed = slide_notes(
            candidate.slides[TARGET_SLIDE_INDEX]
        ) == slide_notes(source.slides[TARGET_SLIDE_INDEX])
    criteria.append(
        result_item(
            "R005",
            target_nonchart_content_passed
            and target_nonchart_style_passed
            and target_notes_passed,
            f"第 3 页非图表内容保持={target_nonchart_content_passed}，"
            f"样式/几何保持={target_nonchart_style_passed}，"
            f"备注保持={target_notes_passed}",
        )
    )

    layout_passed, layout_evidence = chart_layout_checks(
        target_chart,
        candidate.slide_width,
        candidate.slide_height,
    )
    criteria.append(
        result_item(
            "R006",
            layout_passed,
            f"图表布局={json.dumps(layout_evidence, ensure_ascii=False)}",
        )
    )

    if target_chart is not None:
        labels_passed, labels_evidence = chart_label_evidence(
            target_chart,
            len(chart_categories(target_chart.chart)),
        )
    else:
        labels_passed, labels_evidence = False, {"reason": "缺失目标图表"}
    criteria.append(
        result_item(
            "R007",
            labels_passed,
            f"图例/数据标签={json.dumps(labels_evidence, ensure_ascii=False)}",
        )
    )

    try:
        source_media = media_hashes(source_path)
        candidate_media = media_hashes(path)
        source_media_preserved = all(
            candidate_media[value] >= count for value, count in source_media.items()
        )
    except Exception:
        source_media = Counter()
        candidate_media = Counter()
        source_media_preserved = False
    flat_pictures = full_slide_pictures(candidate)
    native_passed = (
        package_passed
        and pie_passed
        and not flat_pictures
        and source_media_preserved
    )
    criteria.append(
        result_item(
            "R009",
            native_passed,
            f"整页图片={flat_pictures}，原生饼图={pie_passed}，"
            f"原媒体保留={source_media_preserved}，"
            f"媒体数={sum(source_media.values())}→{sum(candidate_media.values())}",
        )
    )

    artifact_files = sorted(
        str(item.relative_to(artifact_dir))
        for item in artifact_dir.rglob("*")
        if item.is_file()
    )
    all_text = "\n".join(
        shape.text
        for slide in candidate.slides
        for shape in iter_shapes(slide.shapes)
        if getattr(shape, "has_text_frame", False)
    )
    all_notes = "\n".join(
        text for slide in candidate.slides for text in slide_notes(slide)
    )
    lowered = f"{all_text}\n{all_notes}".lower()
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker.lower() in lowered]
    hygiene_passed = artifact_files == [OUTPUT_NAME] and not markers
    criteria.append(
        result_item(
            "R011",
            hygiene_passed,
            f"输出文件={artifact_files}，可疑标记={markers}",
        )
    )

    pre_render_ready = all(
        next(
            item for item in criteria if item["criterion_id"] == identifier
        )["raw"]["passed"]
        for identifier in ["R001", "R002", "R003", "R004", "R005", "R007", "R009"]
    )
    visual_preservation_passed = False
    non_target_visual_min = 0.0
    target_outside_similarity = 0.0
    render_passed = False
    crop_metrics = {"stddev": 0.0, "nonwhite_ratio": 0.0}
    render_evidence = "前置结构、数据或保持检查失败，未执行渲染"

    if pre_render_ready and source_chart is not None and target_chart is not None:
        with tempfile.TemporaryDirectory(prefix="pptarena029-judge-") as temporary:
            root = Path(temporary)
            source_images, source_pages, source_note = render_presentation(
                source_path, root / "source"
            )
            candidate_images, candidate_pages, candidate_note = render_presentation(
                path, root / "candidate"
            )
            render_evidence = f"源[{source_note}]；候选[{candidate_note}]"
            if len(source_images) == 9 and len(candidate_images) == 9:
                non_target_scores = [
                    image_similarity(source_images[index], candidate_images[index])
                    for index in range(9)
                    if index != TARGET_SLIDE_INDEX
                ]
                non_target_visual_min = min(non_target_scores)
                target_outside_similarity = masked_image_similarity(
                    source_images[TARGET_SLIDE_INDEX],
                    candidate_images[TARGET_SLIDE_INDEX],
                    [geometry(source_chart), geometry(target_chart)],
                    (candidate.slide_width, candidate.slide_height),
                )
                crop_metrics = chart_crop_metrics(
                    candidate_images[TARGET_SLIDE_INDEX],
                    target_chart,
                    (candidate.slide_width, candidate.slide_height),
                )
            visual_preservation_passed = (
                source_pages == 9
                and candidate_pages == 9
                and len(source_images) == 9
                and len(candidate_images) == 9
                and non_target_visual_min >= 0.9985
                and target_outside_similarity >= 0.9985
            )
            render_passed = (
                candidate_pages == 9
                and len(candidate_images) == 9
                and crop_metrics.get("stddev", 0.0) >= 12.0
                and crop_metrics.get("nonwhite_ratio", 0.0) >= 0.03
                and layout_evidence.get("inside", False)
            )

    criteria.append(
        result_item(
            "R008",
            visual_preservation_passed,
            f"非目标页最低像素相似度={non_target_visual_min:.6f}，"
            f"第 3 页图表区外相似度={target_outside_similarity:.6f}；"
            f"{render_evidence}",
        )
    )
    criteria.append(
        result_item(
            "R010",
            render_passed,
            f"饼图区灰度标准差={crop_metrics.get('stddev', 0.0):.4f}，"
            f"非近白像素比例={crop_metrics.get('nonwhite_ratio', 0.0):.4f}，"
            f"图表在页内={layout_evidence.get('inside', False)}；{render_evidence}",
        )
    )

    severe = (
        not pre_render_ready
        or not visual_preservation_passed
        or not render_passed
        or bool(flat_pictures)
        or bool(markers)
    )
    criteria.append(result_item("R012", not severe, f"严重错误={severe}"))
    order = {identifier: index for index, identifier in enumerate(SPECS)}
    criteria.sort(key=lambda item: order[item["criterion_id"]])
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    output = args.artifact_dir / OUTPUT_NAME
    if output.is_file():
        evaluation = evaluate(output, args.artifact_dir)
    else:
        evaluation = {
            "task_id": TASK_ID,
            "rubric_version": RUBRIC_VERSION,
            "judge_version": JUDGE_VERSION,
            "artifact_sha256": [],
            "criterion_results": [],
            "reward": 0.0,
            "passed": False,
            "hurdle_failures": ["missing_output"],
        }
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(
        json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
