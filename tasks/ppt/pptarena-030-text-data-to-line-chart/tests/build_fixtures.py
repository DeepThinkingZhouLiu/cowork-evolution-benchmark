#!/usr/bin/env python3
"""只为 PPTArena Case 30 构造可重放的合法变体、负例和中等例。"""

from __future__ import annotations

import argparse
import importlib.util
import shutil
import sys
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches


OUTPUT_NAME = "LineGraph_GroundTruthA.pptx"
TARGET_SLIDE_INDEX = 7


def load_builder(path: Path):
    spec = importlib.util.spec_from_file_location("pptarena030_builder", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法加载本题 Golden 构造器")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def output_path(root: Path, case: str) -> Path:
    target = root / case / OUTPUT_NAME
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def exact_chart(presentation: Presentation):
    charts = [
        shape
        for shape in presentation.slides[TARGET_SLIDE_INDEX].shapes
        if shape.shape_type == MSO_SHAPE_TYPE.CHART
    ]
    if len(charts) != 1:
        raise ValueError(f"第 8 页应有唯一图表，实际为 {len(charts)}")
    return charts[0]


def valid_alternate(source: Path, builder, output: Path) -> None:
    """从冻结输入独立重建不依赖 Golden 颜色与坐标的合法答案。"""
    presentation = Presentation(source)
    values = builder.parse_source_table(presentation)
    builder.remove_source_table(presentation)
    palette = builder.Palette(nhopi="6A3D9A", asian="E66101", white="4D9221")
    builder.add_line_chart(
        presentation,
        values,
        palette,
        left=0.82,
        top=1.92,
        width=6.95,
        height=4.48,
    )
    builder.add_custom_legend(presentation, palette, left=8.02, top=2.95, row_gap=0.62)
    presentation.save(output)


def wrong_chart_type(source: Path, builder, output: Path) -> None:
    presentation = Presentation(source)
    values = builder.parse_source_table(presentation)
    builder.remove_source_table(presentation)
    data = ChartData()
    data.categories = builder.YEARS
    for name in ("NHOPI", "Asian", "White"):
        data.add_series(name, values[name], number_format="0%")
    presentation.slides[TARGET_SLIDE_INDEX].shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.72),
        Inches(1.88),
        Inches(7.15),
        Inches(4.62),
        data,
    )
    builder.add_custom_legend(presentation, builder.Palette())
    presentation.save(output)


def mutate_first_chart_value(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    chart = exact_chart(presentation).chart
    nodes = chart._element.xpath("./c:chart/c:plotArea/c:lineChart/c:ser[1]/c:val//c:pt[@idx='0']/c:v")
    if len(nodes) != 1:
        raise ValueError("无法定位第一系列第一个值")
    nodes[0].text = "0.35"
    presentation.save(output)


def mutate_first_year(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    chart = exact_chart(presentation).chart
    nodes = chart._element.xpath("./c:chart/c:plotArea/c:lineChart/c:ser[1]/c:cat//c:pt[@idx='0']/c:v")
    if len(nodes) != 1:
        raise ValueError("无法定位第一个年份")
    nodes[0].text = "2012"
    presentation.save(output)


def no_markers(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    for series in exact_chart(presentation).chart.series:
        series.marker.style = XL_MARKER_STYLE.NONE
    presentation.save(output)


def default_legend(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    exact_chart(presentation).chart.has_legend = True
    presentation.save(output)


def missing_custom_legend(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    slide = presentation.slides[TARGET_SLIDE_INDEX]
    for shape in list(slide.shapes):
        if shape.name.startswith("Custom legend"):
            shape._element.getparent().remove(shape._element)
    presentation.save(output)


def mismatched_legend(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    slide = presentation.slides[TARGET_SLIDE_INDEX]
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and shape.text.strip() == "NHOPI":
            shape.text = "Other"
            presentation.save(output)
            return
    raise ValueError("无法定位 NHOPI 图例标签")


def wrong_axis(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    chart = exact_chart(presentation).chart
    chart.value_axis.minimum_scale = 0.05
    chart.value_axis.maximum_scale = 0.40
    chart.value_axis.tick_labels.number_format = "0.00"
    presentation.save(output)


def retained_source_text(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    box = presentation.slides[TARGET_SLIDE_INDEX].shapes.add_textbox(
        Inches(2.5), Inches(2.8), Inches(5.8), Inches(1.9)
    )
    box.text = (
        "Year NHOPI (%) Asian (%) White (%)\n"
        "2013 18 16 12\n2014 14 11 10\n2015 10 8 8\n2016 11 7 7\n2017 11 7 7"
    )
    presentation.save(output)


def changed_slide8_title(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    slide = presentation.slides[TARGET_SLIDE_INDEX]
    for shape in slide.shapes:
        if shape.shape_id == 3 and getattr(shape, "has_text_frame", False):
            shape.text = "Uninsured rates rose under the ACA."
            presentation.save(output)
            return
    raise ValueError("无法定位第 8 页标题")


def changed_other_slide(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    for shape in presentation.slides[0].shapes:
        if getattr(shape, "has_text_frame", False) and "Health and Health Care" in shape.text:
            shape.text = "Changed benchmark answer"
            presentation.save(output)
            return
    raise ValueError("无法定位第 1 页标题")


def duplicate_chart(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    data = ChartData()
    data.categories = ["A", "B"]
    data.add_series("extra", [1, 2])
    presentation.slides[TARGET_SLIDE_INDEX].shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        Inches(7.8),
        Inches(5.0),
        Inches(1.5),
        Inches(1.0),
        data,
    )
    presentation.save(output)


def flattened_slide8(golden: Path, slide_image: Path, output: Path) -> None:
    presentation = Presentation(golden)
    slide = presentation.slides[TARGET_SLIDE_INDEX]
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    slide.shapes.add_picture(
        str(slide_image),
        0,
        0,
        width=presentation.slide_width,
        height=presentation.slide_height,
    )
    presentation.save(output)


def hardcoded_marker(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    box = presentation.slides[TARGET_SLIDE_INDEX].shapes.add_textbox(
        Inches(0.10), Inches(7.10), Inches(2.50), Inches(0.20)
    )
    box.text = "solution/Golden"
    presentation.save(output)


def undersized_chart(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    chart = exact_chart(presentation)
    chart.left = Inches(2.0)
    chart.top = Inches(2.25)
    chart.width = Inches(4.8)
    chart.height = Inches(3.2)
    presentation.save(output)


def extra_slide(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    presentation.slides.add_slide(presentation.slide_layouts[0])
    presentation.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--builder", type=Path, required=True)
    parser.add_argument("--slide8-image", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()
    builder = load_builder(args.builder)

    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))
    valid_alternate(args.source, builder, output_path(args.outdir, "valid-alternate"))
    wrong_chart_type(args.source, builder, output_path(args.outdir, "wrong-chart-type"))
    mutate_first_chart_value(args.golden, output_path(args.outdir, "wrong-value"))
    mutate_first_year(args.golden, output_path(args.outdir, "wrong-year"))
    no_markers(args.golden, output_path(args.outdir, "no-markers"))
    default_legend(args.golden, output_path(args.outdir, "default-legend"))
    missing_custom_legend(args.golden, output_path(args.outdir, "missing-custom-legend"))
    mismatched_legend(args.golden, output_path(args.outdir, "mismatched-legend"))
    wrong_axis(args.golden, output_path(args.outdir, "wrong-axis"))
    retained_source_text(args.golden, output_path(args.outdir, "retained-source-text"))
    changed_slide8_title(args.golden, output_path(args.outdir, "changed-slide8-title"))
    changed_other_slide(args.golden, output_path(args.outdir, "changed-other-slide"))
    duplicate_chart(args.golden, output_path(args.outdir, "duplicate-chart"))
    flattened_slide8(args.golden, args.slide8_image, output_path(args.outdir, "flattened-slide8"))
    hardcoded_marker(args.golden, output_path(args.outdir, "hardcoded-marker"))
    undersized_chart(args.golden, output_path(args.outdir, "undersized-chart"))
    extra_slide(args.golden, output_path(args.outdir, "extra-slide"))


if __name__ == "__main__":
    main()
