#!/usr/bin/env python3
"""PPTArena Case 30 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import statistics
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from PIL import Image, ImageChops, ImageDraw, ImageStat
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "LineGraph_GroundTruthA.pptx"
INPUT_NAME = "LineGraph_TestA.pptx"
INPUT_SHA256 = "72fb98cc839fc54dc4b07dcdb4bf7eb8c6552ac6549d00c0010e3ab021c1f3e6"
TASK_ID = "cowork-evo/pptarena-030-text-data-to-line-chart"
RUBRIC_VERSION = "pptarena-case-030-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-030-deterministic-v0.1"
EMU_PER_INCH = 914_400
TARGET_SLIDE_INDEX = 7
SOURCE_TABLE_SHAPE_ID = 18
PRESERVED_TARGET_SHAPE_IDS = (3, 4, 12)
EXPECTED_NAMES = ("NHOPI", "Asian", "White")
EXPECTED_YEARS = ("2013", "2014", "2015", "2016", "2017")
ALLOWED_LINE_TYPES = {XL_CHART_TYPE.LINE_MARKERS}
SUSPICIOUS_MARKERS = [
    "solution/golden",
    "/root/",
    "judge-result",
    "artifact_dir",
    "pptarena-case-030",
]

SPECS = {
    "R001": ("指定 PPTX 可打开，页面尺寸不变且总页数为 9", 5, "hurdle"),
    "R002": ("第 8 页恰有一张带标记折线图且隐藏默认图例", 14, "hurdle"),
    "R003": ("三系列、五年份和 15 个百分比与冻结输入一致", 18, "hurdle"),
    "R004": ("源数据文本框已删除，第 8 页非目标内容保持", 10, "hurdle"),
    "R005": ("除目标页外的内容、样式、顺序和备注保持", 12, "hurdle"),
    "R006": ("百分比坐标轴正确且三系列视觉可区分", 10, "hurdle"),
    "R007": ("右侧三项原生自定义图例与系列正确对应", 10, "hurdle"),
    "R008": ("图表和自定义图例位于主内容区且无重叠", 7, "positive"),
    "R009": ("非目标页与第 8 页目标区外的渲染效果保持", 6, "hurdle"),
    "R010": ("图表与图例原生可编辑，无压平、宏或媒体丢失", 5, "hurdle"),
    "R011": ("9 页可渲染，目标图表实际可见且非空", 2, "positive"),
    "R012": ("输出目录干净且无答案路径或未授权文字", 1, "positive"),
    "R013": ("页数、图表、数据、图例、原页或可编辑性无严重错误", -20, "penalty"),
}


def normalize_visible(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def locate_source() -> Path | None:
    candidates = [
        Path("/app/input_files") / INPUT_NAME,
        Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME,
    ]
    return next((path for path in candidates if path.is_file()), None)


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def geometry(shape) -> tuple[int, int, int, int]:
    return (int(shape.left), int(shape.top), int(shape.width), int(shape.height))


def rounded_geometry(shape) -> tuple[int, int, int, int]:
    return tuple(round(value / 1000) * 1000 for value in geometry(shape))


def placeholder_type(shape) -> str:
    if not getattr(shape, "is_placeholder", False):
        return ""
    return str(shape.placeholder_format.type)


def color_signature(color_format) -> tuple:
    try:
        color_type = str(color_format.type)
    except Exception:
        color_type = ""
    try:
        rgb = str(color_format.rgb) if color_format.rgb is not None else ""
    except Exception:
        rgb = ""
    try:
        theme = str(color_format.theme_color) if color_format.theme_color is not None else ""
    except Exception:
        theme = ""
    try:
        brightness = round(float(color_format.brightness), 5)
    except Exception:
        brightness = None
    return (color_type, rgb, theme, brightness)


def line_color_signature(line_format) -> tuple:
    try:
        return color_signature(line_format.color)
    except Exception:
        return ("", "", "", None)


def fill_color_signature(fill_format) -> tuple:
    try:
        return color_signature(fill_format.fore_color)
    except Exception:
        return ("", "", "", None)


def run_style_signature(font) -> tuple:
    return (
        font.name,
        round(font.size.pt, 4) if font.size else None,
        bool(font.bold),
        bool(font.italic),
        bool(font.underline),
        int(font.language_id) if font.language_id is not None else None,
        color_signature(font.color),
    )


def text_format_signature(shape) -> tuple:
    text_frame = shape.text_frame
    paragraphs = []
    for paragraph in text_frame.paragraphs:
        styles = {
            run_style_signature(run.font)
            for run in paragraph.runs
            if run.text.strip()
        }
        paragraphs.append(
            (
                normalize_visible(paragraph.text),
                paragraph.level,
                str(paragraph.alignment),
                str(paragraph.line_spacing),
                str(paragraph.space_before),
                str(paragraph.space_after),
                tuple(sorted(styles, key=repr)),
            )
        )
    return (
        int(shape.shape_type),
        placeholder_type(shape),
        rounded_geometry(shape),
        (
            text_frame.margin_left,
            text_frame.margin_right,
            text_frame.margin_top,
            text_frame.margin_bottom,
            str(text_frame.vertical_anchor),
            text_frame.word_wrap,
            str(text_frame.auto_size),
        ),
        tuple(paragraphs),
    )


def chart_categories(chart) -> tuple[str, ...]:
    try:
        return tuple(normalize_visible(str(value)) for value in chart.plots[0].categories)
    except Exception:
        return ()


def numeric_signature(value):
    if value is None:
        return None
    try:
        return round(float(value), 12)
    except Exception:
        return str(value)


def chart_data_signature(chart) -> tuple:
    return (
        str(chart.chart_type),
        chart_categories(chart),
        tuple(
            (
                normalize_visible(series.name or ""),
                tuple(numeric_signature(value) for value in series.values),
            )
            for series in chart.series
        ),
    )


def shape_content_signature(shape) -> tuple:
    text = normalize_visible(shape.text) if getattr(shape, "has_text_frame", False) else ""
    picture_hash = ""
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        picture_hash = hashlib.sha256(shape.image.blob).hexdigest()
    chart = ()
    if shape.shape_type == MSO_SHAPE_TYPE.CHART:
        chart = chart_data_signature(shape.chart)
    children = ()
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        children = tuple(shape_content_signature(child) for child in shape.shapes)
    return (str(shape.shape_type), placeholder_type(shape), text, picture_hash, chart, children)


def shape_style_signature(shape) -> tuple:
    text_style = text_format_signature(shape) if getattr(shape, "has_text_frame", False) else ()
    picture_hash = ""
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        picture_hash = hashlib.sha256(shape.image.blob).hexdigest()
    chart_style = ()
    if shape.shape_type == MSO_SHAPE_TYPE.CHART:
        chart_style = (chart_data_signature(shape.chart), bool(shape.chart.has_legend), shape.chart.chart_style)
    children = ()
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        children = tuple(shape_style_signature(child) for child in shape.shapes)
    return (
        str(shape.shape_type),
        placeholder_type(shape),
        rounded_geometry(shape),
        text_style,
        picture_hash,
        chart_style,
        children,
    )


def slide_content_signature(slide) -> tuple:
    return tuple(shape_content_signature(shape) for shape in slide.shapes)


def slide_style_signature(slide) -> tuple:
    return tuple(shape_style_signature(shape) for shape in slide.shapes)


def slide_notes(slide) -> tuple[str, ...]:
    try:
        return tuple(
            shape.text
            for shape in slide.notes_slide.shapes
            if getattr(shape, "has_text_frame", False)
        )
    except Exception:
        return ()


def media_hashes(path: Path) -> Counter[str]:
    with zipfile.ZipFile(path) as archive:
        return Counter(
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("ppt/media/")
        )


def package_ok(path: Path) -> tuple[bool, str]:
    try:
        with zipfile.ZipFile(path) as archive:
            broken = archive.testzip()
            names = set(archive.namelist())
            required = {
                "[Content_Types].xml",
                "ppt/presentation.xml",
                "ppt/slides/slide1.xml",
                "ppt/slides/slide8.xml",
                "ppt/slides/slide9.xml",
            }
            missing = sorted(required - names)
            macros = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
        return broken is None and not missing and not macros, (
            f"损坏成员={broken}，缺失={missing}，宏={macros}"
        )
    except Exception as exc:
        return False, f"ZIP 检查失败：{exc}"


def chart_shapes(slide) -> list:
    return [shape for shape in iter_shapes(slide.shapes) if shape.shape_type == MSO_SHAPE_TYPE.CHART]


def parse_source_values(source: Presentation) -> tuple[tuple[str, ...], dict[str, tuple[float, ...]]]:
    slide = source.slides[TARGET_SLIDE_INDEX]
    matches = [shape for shape in slide.shapes if shape.shape_id == SOURCE_TABLE_SHAPE_ID]
    if len(matches) != 1 or not matches[0].has_text_frame:
        raise ValueError("无法定位冻结输入数据文本框")
    lines = [normalize_visible(line) for line in matches[0].text.splitlines() if normalize_visible(line)]
    if len(lines) != 6:
        raise ValueError("冻结输入数据文本框行数异常")
    rows = []
    for line in lines[1:]:
        numbers = [int(value) for value in re.findall(r"\d+", line)]
        if len(numbers) != 4:
            raise ValueError(f"无法解析冻结数据行：{line!r}")
        rows.append(numbers)
    years = tuple(str(row[0]) for row in rows)
    values = {
        "NHOPI": tuple(row[1] / 100 for row in rows),
        "Asian": tuple(row[2] / 100 for row in rows),
        "White": tuple(row[3] / 100 for row in rows),
    }
    return years, values


def compare_chart_data(source: Presentation, chart_shape) -> tuple[bool, dict]:
    expected_years, expected_values = parse_source_values(source)
    chart = chart_shape.chart
    actual_years = chart_categories(chart)
    actual_series = {
        normalize_visible(series.name or ""): tuple(
            float(value) if value is not None else None for value in series.values
        )
        for series in chart.series
    }
    mismatches = []
    for name in EXPECTED_NAMES:
        expected = expected_values.get(name, ())
        actual = actual_series.get(name, ())
        if len(expected) != len(actual):
            mismatches.append({"series": name, "reason": "数据点数量不一致"})
            continue
        for year, expected_value, actual_value in zip(expected_years, expected, actual, strict=True):
            if actual_value is None or not math.isclose(
                float(expected_value), float(actual_value), rel_tol=1e-9, abs_tol=1e-9
            ):
                mismatches.append(
                    {"series": name, "year": year, "expected": expected_value, "actual": actual_value}
                )
    passed = (
        expected_years == EXPECTED_YEARS
        and actual_years == expected_years
        and set(actual_series) == set(EXPECTED_NAMES)
        and len(actual_series) == 3
        and not mismatches
    )
    return passed, {
        "expected_years": list(expected_years),
        "actual_years": list(actual_years),
        "expected_series": list(EXPECTED_NAMES),
        "actual_series": sorted(actual_series),
        "mismatches": mismatches,
    }


def marker_is_enabled(series) -> bool:
    try:
        style = series.marker.style
    except Exception:
        return True
    return style is None or "NONE" not in str(style).upper()


def series_visual_signature(series) -> tuple:
    try:
        marker_style = str(series.marker.style)
    except Exception:
        marker_style = ""
    try:
        marker_size = int(series.marker.size) if series.marker.size is not None else None
    except Exception:
        marker_size = None
    return (line_color_signature(series.format.line), marker_style, marker_size)


def axis_and_series_checks(chart_shape) -> tuple[bool, dict]:
    chart = chart_shape.chart
    try:
        minimum = chart.value_axis.minimum_scale
        maximum = chart.value_axis.maximum_scale
        number_format = str(chart.value_axis.tick_labels.number_format or "")
    except Exception:
        minimum, maximum, number_format = None, None, ""
    marker_flags = [marker_is_enabled(series) for series in chart.series]
    visual_signatures = [series_visual_signature(series) for series in chart.series]
    distinct = len(set(visual_signatures)) == len(visual_signatures) == 3
    percentage_format = "%" in number_format
    scale_ok = (
        minimum is not None
        and maximum is not None
        and float(minimum) <= 0.001
        and 0.195 <= float(maximum) <= 0.205
    )
    passed = scale_ok and percentage_format and all(marker_flags) and distinct
    return passed, {
        "minimum_scale": minimum,
        "maximum_scale": maximum,
        "number_format": number_format,
        "markers_enabled": marker_flags,
        "series_visual_signatures": [repr(value) for value in visual_signatures],
        "series_distinct": distinct,
    }


def rect_center_y(shape) -> float:
    return float(shape.top + shape.height / 2)


def legend_checks(slide, chart_shape) -> tuple[bool, dict, list]:
    shapes = list(iter_shapes(slide.shapes))
    labels = {
        name: [
            shape
            for shape in shapes
            if getattr(shape, "has_text_frame", False) and normalize_visible(shape.text) == name
        ]
        for name in EXPECTED_NAMES
    }
    lines = [shape for shape in shapes if shape.shape_type == MSO_SHAPE_TYPE.LINE]
    markers = [
        shape
        for shape in shapes
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
        and shape.width <= int(0.35 * EMU_PER_INCH)
        and shape.height <= int(0.35 * EMU_PER_INCH)
    ]
    chart_right = chart_shape.left + chart_shape.width
    series_by_name = {normalize_visible(series.name or ""): series for series in chart_shape.chart.series}
    evidence = {}
    used_shapes = []
    all_passed = True
    for name in EXPECTED_NAMES:
        if len(labels[name]) != 1:
            evidence[name] = {"label_count": len(labels[name])}
            all_passed = False
            continue
        label = labels[name][0]
        row_lines = [
            shape
            for shape in lines
            if shape.left >= chart_right
            and shape.left + shape.width <= label.left + int(0.05 * EMU_PER_INCH)
            and abs(rect_center_y(shape) - rect_center_y(label)) <= int(0.20 * EMU_PER_INCH)
            and shape.width >= int(0.12 * EMU_PER_INCH)
        ]
        row_markers = [
            shape
            for shape in markers
            if shape.left >= chart_right
            and shape.left + shape.width <= label.left + int(0.08 * EMU_PER_INCH)
            and abs(rect_center_y(shape) - rect_center_y(label)) <= int(0.20 * EMU_PER_INCH)
        ]
        series = series_by_name.get(name)
        expected_color = line_color_signature(series.format.line) if series is not None else ()
        line_colors = [line_color_signature(shape.line) for shape in row_lines]
        marker_colors = [fill_color_signature(shape.fill) for shape in row_markers]
        explicit_expected = bool(expected_color and expected_color[1])
        color_matched = (
            not explicit_expected
            or any(color == expected_color for color in line_colors)
            and any(color == expected_color for color in marker_colors)
        )
        row_passed = (
            label.left >= chart_right
            and len(row_lines) >= 1
            and len(row_markers) >= 1
            and color_matched
        )
        all_passed = all_passed and row_passed
        used_shapes.extend([label, *row_lines, *row_markers])
        evidence[name] = {
            "label_box_in": [round(value / EMU_PER_INCH, 3) for value in geometry(label)],
            "line_count": len(row_lines),
            "marker_count": len(row_markers),
            "expected_color": repr(expected_color),
            "line_colors": [repr(value) for value in line_colors],
            "marker_colors": [repr(value) for value in marker_colors],
            "color_matched": color_matched,
        }
    label_order = [labels[name][0].top for name in EXPECTED_NAMES if len(labels[name]) == 1]
    ordered = len(label_order) == 3 and label_order == sorted(label_order)
    all_passed = all_passed and ordered
    evidence["ordered_top_to_bottom"] = ordered
    return all_passed, evidence, used_shapes


def rectangles_overlap(first, second, tolerance: int = 0) -> bool:
    fl, ft, fw, fh = geometry(first)
    sl, st, sw, sh = geometry(second)
    return not (
        fl + fw <= sl + tolerance
        or sl + sw <= fl + tolerance
        or ft + fh <= st + tolerance
        or st + sh <= ft + tolerance
    )


def layout_checks(chart_shape, legend_shapes: list, slide_width: int, slide_height: int) -> tuple[bool, dict]:
    left, top, width, height = geometry(chart_shape)
    right, bottom = left + width, top + height
    chart_inside = (
        left >= int(0.45 * EMU_PER_INCH)
        and top >= int(1.70 * EMU_PER_INCH)
        and right <= int(8.15 * EMU_PER_INCH)
        and bottom <= int(6.60 * EMU_PER_INCH)
    )
    chart_large = width >= int(5.50 * EMU_PER_INCH) and height >= int(3.70 * EMU_PER_INCH)
    legend_inside = bool(legend_shapes) and all(
        shape.left >= right
        and shape.left >= int(7.70 * EMU_PER_INCH)
        and shape.left + shape.width <= slide_width - int(0.05 * EMU_PER_INCH)
        and shape.top >= int(1.75 * EMU_PER_INCH)
        and shape.top + shape.height <= int(6.45 * EMU_PER_INCH)
        for shape in legend_shapes
    )
    overlap = any(rectangles_overlap(chart_shape, shape) for shape in legend_shapes)
    return chart_inside and chart_large and legend_inside and not overlap, {
        "chart_box_in": [round(value / EMU_PER_INCH, 3) for value in geometry(chart_shape)],
        "chart_inside": chart_inside,
        "chart_large_enough": chart_large,
        "legend_shape_count": len(legend_shapes),
        "legend_inside": legend_inside,
        "chart_legend_overlap": overlap,
        "slide_size_in": [round(slide_width / EMU_PER_INCH, 3), round(slide_height / EMU_PER_INCH, 3)],
    }


def full_slide_pictures(presentation: Presentation) -> list[str]:
    values = []
    page_area = presentation.slide_width * presentation.slide_height
    for slide_index, slide in enumerate(presentation.slides, start=1):
        for shape in iter_shapes(slide.shapes):
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            coverage = shape.width * shape.height / page_area
            if coverage >= 0.85:
                values.append(f"slide{slide_index}:{shape.name}:{coverage:.3f}")
    return values


def render_presentation(path: Path, outdir: Path) -> tuple[list[Path], int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "lo-profile"
    profile.mkdir()
    completed = subprocess.run(
        [
            "soffice",
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(outdir),
            str(path),
        ],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return [], 0, f"转 PDF 失败：{completed.stdout[-500:]}"
    info = subprocess.run(
        ["pdfinfo", str(pdf)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=30,
        check=False,
    ).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    pages = int(match.group(1)) if match else 0
    prefix = outdir / "slide"
    raster = subprocess.run(
        ["pdftoppm", "-jpeg", "-r", "90", str(pdf), str(prefix)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    images = sorted(outdir.glob("slide-*.jpg")) if raster.returncode == 0 else []
    return images, pages, (
        f"页数={pages}，渲染图={len(images)}，"
        f"soffice={completed.returncode}，pdftoppm={raster.returncode}"
    )


def image_similarity(first: Path, second: Path) -> float:
    with Image.open(first) as left_image, Image.open(second) as right_image:
        left, right = left_image.convert("RGB"), right_image.convert("RGB")
        if left.size != right.size:
            return 0.0
        histogram = ImageChops.difference(left, right).histogram()
        squared = sum(count * ((index % 256) ** 2) for index, count in enumerate(histogram))
        pixels = left.width * left.height
        rms = math.sqrt(squared / (pixels * 3)) if pixels else 255.0
        return max(0.0, 1.0 - rms / 255.0)


def emu_rect_to_pixels(rectangle, image_size, slide_size, padding: int = 0):
    image_width, image_height = image_size
    slide_width, slide_height = slide_size
    left, top, width, height = rectangle
    return (
        max(0, math.floor(left * image_width / slide_width) - padding),
        max(0, math.floor(top * image_height / slide_height) - padding),
        min(image_width, math.ceil((left + width) * image_width / slide_width) + padding),
        min(image_height, math.ceil((top + height) * image_height / slide_height) + padding),
    )


def masked_image_similarity(first: Path, second: Path, ignored_rectangles, slide_size) -> float:
    with Image.open(first) as first_image, Image.open(second) as second_image:
        left, right = first_image.convert("RGB"), second_image.convert("RGB")
        if left.size != right.size:
            return 0.0
        difference = ImageChops.difference(left, right)
        mask = Image.new("L", left.size, 255)
        draw = ImageDraw.Draw(mask)
        for rectangle in ignored_rectangles:
            draw.rectangle(emu_rect_to_pixels(rectangle, left.size, slide_size, padding=10), fill=0)
        stats = ImageStat.Stat(difference, mask=mask)
        rms = statistics.mean(stats.rms) if stats.rms else 255.0
        return max(0.0, 1.0 - rms / 255.0)


def target_crop_metrics(image_path: Path, slide_size) -> dict:
    target_rectangle = (
        int(0.45 * EMU_PER_INCH),
        int(1.72 * EMU_PER_INCH),
        int(9.35 * EMU_PER_INCH),
        int(4.85 * EMU_PER_INCH),
    )
    with Image.open(image_path) as image:
        rgb = image.convert("RGB")
        crop = rgb.crop(emu_rect_to_pixels(target_rectangle, rgb.size, slide_size))
        stddev = statistics.mean(ImageStat.Stat(crop.convert("L")).stddev)
        histogram = crop.convert("L").histogram()
        nonwhite = sum(histogram[:245]) / (crop.width * crop.height)
        return {"stddev": float(stddev), "nonwhite_ratio": float(nonwhite)}


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {
        "criterion_id": identifier,
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
        "raw": {
            "description": description,
            "weight": weight,
            "type": kind,
            "passed": bool(passed),
        },
    }


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [
        item["criterion_id"]
        for item in criteria
        if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]
    ]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(
        item["raw"]["weight"]
        for item in criteria
        if item["raw"]["weight"] > 0 and item["raw"]["passed"]
    )
    penalty = sum(
        abs(item["raw"]["weight"])
        for item in criteria
        if item["raw"]["weight"] < 0 and not item["raw"]["passed"]
    )
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [sha256(path)] if path.is_file() else [],
        "criterion_results": criteria,
        "reward": reward,
        "passed": reward >= 0.8 and not hurdles,
        "hurdle_failures": hurdles,
    }


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        criteria = [result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配") for identifier in SPECS]
        return make_result(path, criteria)
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        criteria = [result_item(identifier, False, f"PPTX 打开失败：{exc}") for identifier in SPECS]
        return make_result(path, criteria)

    criteria = []
    package_passed, package_evidence = package_ok(path)
    size_passed = (
        len(source.slides) == 9
        and len(candidate.slides) == 9
        and candidate.slide_width == source.slide_width
        and candidate.slide_height == source.slide_height
    )
    criteria.append(
        result_item(
            "R001",
            package_passed and size_passed,
            f"{package_evidence}，输入页数={len(source.slides)}，候选页数={len(candidate.slides)}，"
            f"页面={candidate.slide_width / EMU_PER_INCH:.4f}×{candidate.slide_height / EMU_PER_INCH:.4f}英寸",
        )
    )

    target_slide = candidate.slides[TARGET_SLIDE_INDEX] if len(candidate.slides) > TARGET_SLIDE_INDEX else None
    target_charts = chart_shapes(target_slide) if target_slide is not None else []
    target_chart = target_charts[0] if len(target_charts) == 1 else None
    line_passed = (
        target_chart is not None
        and target_chart.chart.chart_type in ALLOWED_LINE_TYPES
        and not target_chart.chart.has_legend
        and all(marker_is_enabled(series) for series in target_chart.chart.series)
    )
    criteria.append(
        result_item(
            "R002",
            line_passed,
            f"第 8 页图表数={len(target_charts)}，类型="
            f"{str(target_chart.chart.chart_type) if target_chart is not None else '缺失'}，"
            f"默认图例={bool(target_chart.chart.has_legend) if target_chart is not None else '缺失'}",
        )
    )

    if target_chart is not None:
        try:
            data_passed, data_evidence = compare_chart_data(source, target_chart)
        except Exception as exc:
            data_passed, data_evidence = False, {"reason": str(exc)}
    else:
        data_passed, data_evidence = False, {"reason": "缺失唯一目标图表"}
    criteria.append(result_item("R003", data_passed, json.dumps(data_evidence, ensure_ascii=False)))

    preserved_checks = {}
    source_text_removed = False
    unauthorized_text = []
    target_notes_passed = False
    if target_slide is not None:
        source_target = source.slides[TARGET_SLIDE_INDEX]
        for shape_id in PRESERVED_TARGET_SHAPE_IDS:
            source_matches = [shape for shape in source_target.shapes if shape.shape_id == shape_id]
            candidate_matches = [shape for shape in target_slide.shapes if shape.shape_id == shape_id]
            preserved_checks[str(shape_id)] = (
                len(source_matches) == 1
                and len(candidate_matches) == 1
                and shape_content_signature(source_matches[0]) == shape_content_signature(candidate_matches[0])
                and shape_style_signature(source_matches[0]) == shape_style_signature(candidate_matches[0])
            )
        all_target_text = [
            normalize_visible(shape.text)
            for shape in iter_shapes(target_slide.shapes)
            if getattr(shape, "has_text_frame", False)
        ]
        source_text_removed = not any(
            "NHOPI (%)" in text and "White (%)" in text for text in all_target_text
        )
        preserved_text = {
            normalize_visible(shape.text)
            for shape in source_target.shapes
            if shape.shape_id in PRESERVED_TARGET_SHAPE_IDS and getattr(shape, "has_text_frame", False)
        }
        unauthorized_text = [
            text
            for text in all_target_text
            if text and text not in preserved_text and text not in EXPECTED_NAMES
        ]
        target_notes_passed = slide_notes(target_slide) == slide_notes(source_target)
    target_preserved_passed = (
        all(preserved_checks.values())
        and len(preserved_checks) == len(PRESERVED_TARGET_SHAPE_IDS)
        and source_text_removed
        and not unauthorized_text
        and target_notes_passed
    )
    criteria.append(
        result_item(
            "R004",
            target_preserved_passed,
            f"保留对象={preserved_checks}，源数据文本已删除={source_text_removed}，"
            f"未授权新增文字={unauthorized_text}，备注保持={target_notes_passed}",
        )
    )

    content_mismatches, style_mismatches, notes_mismatches = [], [], []
    if len(candidate.slides) == 9:
        for index in range(9):
            if index == TARGET_SLIDE_INDEX:
                continue
            if slide_content_signature(candidate.slides[index]) != slide_content_signature(source.slides[index]):
                content_mismatches.append(index + 1)
            if slide_style_signature(candidate.slides[index]) != slide_style_signature(source.slides[index]):
                style_mismatches.append(index + 1)
            if slide_notes(candidate.slides[index]) != slide_notes(source.slides[index]):
                notes_mismatches.append(index + 1)
    else:
        content_mismatches = [1, 2, 3, 4, 5, 6, 7, 9]
        style_mismatches = content_mismatches[:]
        notes_mismatches = content_mismatches[:]
    non_target_preserved = not content_mismatches and not style_mismatches and not notes_mismatches
    criteria.append(
        result_item(
            "R005",
            non_target_preserved,
            f"非目标页内容变化={content_mismatches}，样式/几何变化={style_mismatches}，"
            f"备注变化={notes_mismatches}",
        )
    )

    if target_chart is not None:
        axis_passed, axis_evidence = axis_and_series_checks(target_chart)
    else:
        axis_passed, axis_evidence = False, {"reason": "缺失目标图表"}
    criteria.append(result_item("R006", axis_passed, json.dumps(axis_evidence, ensure_ascii=False)))

    if target_slide is not None and target_chart is not None:
        legend_passed, legend_evidence, legend_shapes = legend_checks(target_slide, target_chart)
    else:
        legend_passed, legend_evidence, legend_shapes = False, {"reason": "缺失目标页或图表"}, []
    criteria.append(result_item("R007", legend_passed, json.dumps(legend_evidence, ensure_ascii=False)))

    if target_chart is not None:
        layout_passed, layout_evidence = layout_checks(
            target_chart, legend_shapes, candidate.slide_width, candidate.slide_height
        )
    else:
        layout_passed, layout_evidence = False, {"reason": "缺失目标图表"}
    criteria.append(result_item("R008", layout_passed, json.dumps(layout_evidence, ensure_ascii=False)))

    try:
        source_media = media_hashes(source_path)
        candidate_media = media_hashes(path)
        source_media_preserved = all(candidate_media[value] >= count for value, count in source_media.items())
    except Exception:
        source_media, candidate_media, source_media_preserved = Counter(), Counter(), False
    flat_pictures = full_slide_pictures(candidate)
    native_passed = (
        package_passed
        and line_passed
        and legend_passed
        and not flat_pictures
        and source_media_preserved
    )
    criteria.append(
        result_item(
            "R010",
            native_passed,
            f"整页图片={flat_pictures}，原生折线图={line_passed}，原生自定义图例={legend_passed}，"
            f"原媒体保留={source_media_preserved}，媒体数={sum(source_media.values())}→{sum(candidate_media.values())}",
        )
    )

    artifact_files = sorted(
        str(item.relative_to(artifact_dir)) for item in artifact_dir.rglob("*") if item.is_file()
    )
    all_text = "\n".join(
        shape.text
        for slide in candidate.slides
        for shape in iter_shapes(slide.shapes)
        if getattr(shape, "has_text_frame", False)
    )
    all_notes = "\n".join(text for slide in candidate.slides for text in slide_notes(slide))
    lowered = f"{all_text}\n{all_notes}".lower()
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker.lower() in lowered]
    hygiene_passed = artifact_files == [OUTPUT_NAME] and not markers
    criteria.append(result_item("R012", hygiene_passed, f"输出文件={artifact_files}，可疑标记={markers}"))

    pre_render_ready = all(
        next(item for item in criteria if item["criterion_id"] == identifier)["raw"]["passed"]
        for identifier in ["R001", "R002", "R003", "R004", "R005", "R006", "R007", "R010"]
    )
    visual_preservation_passed = False
    non_target_visual_min = 0.0
    target_outside_similarity = 0.0
    render_passed = False
    crop_metrics = {"stddev": 0.0, "nonwhite_ratio": 0.0}
    render_evidence = "前置结构、数据或保持检查失败，未执行渲染"
    if pre_render_ready:
        with tempfile.TemporaryDirectory(prefix="pptarena030-judge-") as temporary:
            root = Path(temporary)
            source_images, source_pages, source_note = render_presentation(source_path, root / "source")
            candidate_images, candidate_pages, candidate_note = render_presentation(path, root / "candidate")
            render_evidence = f"源[{source_note}]；候选[{candidate_note}]"
            if len(source_images) == 9 and len(candidate_images) == 9:
                non_target_scores = [
                    image_similarity(source_images[index], candidate_images[index])
                    for index in range(9)
                    if index != TARGET_SLIDE_INDEX
                ]
                non_target_visual_min = min(non_target_scores)
                ignored = [
                    (
                        int(0.45 * EMU_PER_INCH),
                        int(1.72 * EMU_PER_INCH),
                        int(9.35 * EMU_PER_INCH),
                        int(4.85 * EMU_PER_INCH),
                    )
                ]
                target_outside_similarity = masked_image_similarity(
                    source_images[TARGET_SLIDE_INDEX],
                    candidate_images[TARGET_SLIDE_INDEX],
                    ignored,
                    (candidate.slide_width, candidate.slide_height),
                )
                crop_metrics = target_crop_metrics(
                    candidate_images[TARGET_SLIDE_INDEX],
                    (candidate.slide_width, candidate.slide_height),
                )
            visual_preservation_passed = (
                source_pages == 9
                and candidate_pages == 9
                and len(source_images) == 9
                and len(candidate_images) == 9
                and non_target_visual_min >= 0.9985
                and target_outside_similarity >= 0.9985
            )
            render_passed = (
                candidate_pages == 9
                and len(candidate_images) == 9
                and crop_metrics.get("stddev", 0.0) >= 12.0
                and crop_metrics.get("nonwhite_ratio", 0.0) >= 0.025
            )

    criteria.append(
        result_item(
            "R009",
            visual_preservation_passed,
            f"非目标页最低像素相似度={non_target_visual_min:.6f}，"
            f"第 8 页目标区外相似度={target_outside_similarity:.6f}；{render_evidence}",
        )
    )
    criteria.append(
        result_item(
            "R011",
            render_passed,
            f"目标区灰度标准差={crop_metrics.get('stddev', 0.0):.4f}，"
            f"非近白像素比例={crop_metrics.get('nonwhite_ratio', 0.0):.4f}；{render_evidence}",
        )
    )

    severe = (
        not pre_render_ready
        or not visual_preservation_passed
        or not render_passed
        or bool(flat_pictures)
        or bool(markers)
    )
    criteria.append(result_item("R013", not severe, f"严重错误={severe}"))
    order = {identifier: index for index, identifier in enumerate(SPECS)}
    criteria.sort(key=lambda item: order[item["criterion_id"]])
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    output = args.artifact_dir / OUTPUT_NAME
    if output.is_file():
        evaluation = evaluate(output, args.artifact_dir)
    else:
        evaluation = {
            "task_id": TASK_ID,
            "rubric_version": RUBRIC_VERSION,
            "judge_version": JUDGE_VERSION,
            "artifact_sha256": [],
            "criterion_results": [],
            "reward": 0.0,
            "passed": False,
            "hurdle_failures": ["missing_output"],
        }
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
