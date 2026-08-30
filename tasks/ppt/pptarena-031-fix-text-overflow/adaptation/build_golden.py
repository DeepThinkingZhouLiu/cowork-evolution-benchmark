#!/usr/bin/env python3
"""PPTArena Case 31 的单题 Golden 构造器。"""

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


INPUT_SHA256 = "183e662d0738577cee4c177dfcefc8620c664538e239b818aab47ee319b9a46e"
TEXT_SHA256 = "416c20667310582f95adada698b9c970cb54079e7023398972794277bad637c8"
TARGET_SHAPE_ID = 2
SOURCE_GEOMETRY_IN = (0.6722, 3.5694, 14.9667, 2.5244)
GOLDEN_WIDTH_IN = 10.653
GOLDEN_HEIGHT_IN = 3.433


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def text_sha256(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def rounded_geometry_in(shape) -> tuple[float, float, float, float]:
    return tuple(round(value / 914_400, 4) for value in (shape.left, shape.top, shape.width, shape.height))


def build(source: Path, output: Path, width_in: float = GOLDEN_WIDTH_IN, height_in: float = GOLDEN_HEIGHT_IN) -> None:
    if not source.is_file() or sha256(source) != INPUT_SHA256:
        raise ValueError("冻结输入缺失或 SHA256 不匹配")
    presentation = Presentation(source)
    if len(presentation.slides) != 1:
        raise ValueError("冻结输入应为 1 页")
    matches = [shape for shape in presentation.slides[0].shapes if shape.shape_id == TARGET_SHAPE_ID]
    if len(matches) != 1 or not matches[0].has_text_frame:
        raise ValueError("无法唯一定位 TextBox 1")
    target = matches[0]
    if rounded_geometry_in(target) != SOURCE_GEOMETRY_IN:
        raise ValueError(f"冻结文本框几何不匹配：{rounded_geometry_in(target)}")
    if text_sha256(target.text) != TEXT_SHA256:
        raise ValueError("冻结文本内容不匹配")
    target.width = Inches(width_in)
    target.height = Inches(height_in)
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.source, args.output)
    print(args.output)


if __name__ == "__main__":
    main()
