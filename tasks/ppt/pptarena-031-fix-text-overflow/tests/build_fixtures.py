#!/usr/bin/env python3
"""只为 PPTArena Case 31 构造合法变体、负例和中等例。"""

from __future__ import annotations

import argparse
import importlib.util
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


OUTPUT_NAME = "FixTextOverFlow_GroundTurthA.pptx"


def load_builder(path: Path):
    spec = importlib.util.spec_from_file_location("pptarena031_builder", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法加载本题 Golden 构造器")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def output_path(root: Path, case: str) -> Path:
    target = root / case / OUTPUT_NAME
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def target_shape(presentation: Presentation):
    matches = [shape for shape in presentation.slides[0].shapes if shape.shape_id == 2]
    if len(matches) != 1:
        raise ValueError("无法唯一定位 TextBox 1")
    return matches[0]


def valid_alternate(source: Path, builder, output: Path) -> None:
    builder.build(source, output, width_in=11.30, height_in=3.20)


def cramped_margin(source: Path, builder, output: Path) -> None:
    builder.build(source, output, width_in=12.55, height_in=3.00)


def moved_top_left(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    target = target_shape(presentation)
    target.left += Inches(0.30)
    presentation.save(output)


def off_page(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    target_shape(presentation).width = Inches(13.0)
    presentation.save(output)


def insufficient_height(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    target_shape(presentation).height = Inches(2.55)
    presentation.save(output)


def altered_text(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    target = target_shape(presentation)
    target.text_frame.paragraphs[0].runs[0].text = target.text_frame.paragraphs[0].runs[0].text.replace(
        "Artificial Intelligence", "Automation"
    )
    presentation.save(output)


def shrunk_font(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    target_shape(presentation).text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    presentation.save(output)


def changed_title(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    for shape in presentation.slides[0].shapes:
        if getattr(shape, "has_text_frame", False) and shape.text.strip() == "Conclusion":
            shape.text = "Summary"
            presentation.save(output)
            return
    raise ValueError("无法定位标题")


def changed_background(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    fill = presentation.slides[0].background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(255, 255, 255)
    presentation.save(output)


def added_shape(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    box = presentation.slides[0].shapes.add_textbox(Inches(11.7), Inches(0.2), Inches(1.2), Inches(0.4))
    box.text = "extra"
    presentation.save(output)


def disabled_wrap(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    target_shape(presentation).text_frame.word_wrap = False
    presentation.save(output)


def flattened(golden: Path, image: Path, output: Path) -> None:
    presentation = Presentation(golden)
    slide = presentation.slides[0]
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    slide.shapes.add_picture(
        str(image),
        0,
        0,
        width=presentation.slide_width,
        height=presentation.slide_height,
    )
    presentation.save(output)


def hardcoded_marker(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    box = presentation.slides[0].shapes.add_textbox(Inches(0.1), Inches(7.1), Inches(2.5), Inches(0.2))
    box.text = "solution/Golden"
    presentation.save(output)


def extra_slide(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    presentation.slides.add_slide(presentation.slide_layouts[0])
    presentation.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--builder", type=Path, required=True)
    parser.add_argument("--slide-image", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()
    builder = load_builder(args.builder)

    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))
    valid_alternate(args.source, builder, output_path(args.outdir, "valid-alternate"))
    cramped_margin(args.source, builder, output_path(args.outdir, "cramped-margin"))
    moved_top_left(args.golden, output_path(args.outdir, "moved-top-left"))
    off_page(args.golden, output_path(args.outdir, "off-page"))
    insufficient_height(args.golden, output_path(args.outdir, "insufficient-height"))
    altered_text(args.golden, output_path(args.outdir, "altered-text"))
    shrunk_font(args.golden, output_path(args.outdir, "shrunk-font"))
    changed_title(args.golden, output_path(args.outdir, "changed-title"))
    changed_background(args.golden, output_path(args.outdir, "changed-background"))
    added_shape(args.golden, output_path(args.outdir, "added-shape"))
    disabled_wrap(args.golden, output_path(args.outdir, "disabled-wrap"))
    flattened(args.golden, args.slide_image, output_path(args.outdir, "flattened"))
    hardcoded_marker(args.golden, output_path(args.outdir, "hardcoded-marker"))
    extra_slide(args.golden, output_path(args.outdir, "extra-slide"))


if __name__ == "__main__":
    main()
