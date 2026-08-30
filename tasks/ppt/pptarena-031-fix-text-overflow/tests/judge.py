#!/usr/bin/env python3
"""PPTArena Case 31 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import statistics
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from PIL import Image, ImageChops, ImageDraw, ImageStat
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "FixTextOverFlow_GroundTurthA.pptx"
INPUT_NAME = "FixTextOverFlow_TestA.pptx"
INPUT_SHA256 = "183e662d0738577cee4c177dfcefc8620c664538e239b818aab47ee319b9a46e"
TASK_ID = "cowork-evo/pptarena-031-fix-text-overflow"
RUBRIC_VERSION = "pptarena-case-031-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-031-deterministic-v0.1"
EMU_PER_INCH = 914_400
POSITION_TOLERANCE = int(0.02 * EMU_PER_INCH)
SUSPICIOUS_MARKERS = ["solution/golden", "/root/", "judge-result", "artifact_dir", "pptarena-case-031"]

SPECS = {
    "R001": ("指定 PPTX 可打开，页面尺寸不变且总页数为 1", 7, "hurdle"),
    "R002": ("主文本内容和文本格式完整保留", 18, "hurdle"),
    "R003": ("主文本框左上角不变且宽高已调整到页内", 15, "hurdle"),
    "R004": ("主文本在实际渲染中全部可见且无溢出", 15, "hurdle"),
    "R005": ("标题、背景、主题、页面和备注保持", 14, "hurdle"),
    "R006": ("页面只保留两个原生可编辑文字对象", 10, "hurdle"),
    "R007": ("文本框边距合理、不与标题重叠且比例自然", 8, "positive"),
    "R008": ("整页可渲染，目标区外保持且文字实际可见", 7, "hurdle"),
    "R009": ("原媒体与包成员保留，无宏和损坏关系", 4, "hurdle"),
    "R010": ("输出目录干净且无答案路径或未授权文字", 2, "positive"),
    "R011": ("文本、格式、边界、标题和可编辑性无严重错误", -20, "penalty"),
}


def normalize_visible(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def normalize_pdf_text(value: str) -> str:
    value = re.sub(r"([A-Za-z])\-\s+([A-Za-z])", r"\1-\2", value or "")
    return normalize_visible(value)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def locate_source() -> Path | None:
    candidates = [
        Path("/app/input_files") / INPUT_NAME,
        Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME,
    ]
    return next((path for path in candidates if path.is_file()), None)


def geometry(shape) -> tuple[int, int, int, int]:
    return (int(shape.left), int(shape.top), int(shape.width), int(shape.height))


def rounded_geometry(shape) -> tuple[int, int, int, int]:
    return tuple(round(value / 1000) * 1000 for value in geometry(shape))


def placeholder_type(shape) -> str:
    if not getattr(shape, "is_placeholder", False):
        return ""
    return str(shape.placeholder_format.type)


def color_signature(color_format) -> tuple:
    try:
        color_type = str(color_format.type)
    except Exception:
        color_type = ""
    try:
        rgb = str(color_format.rgb) if color_format.rgb is not None else ""
    except Exception:
        rgb = ""
    try:
        theme = str(color_format.theme_color) if color_format.theme_color is not None else ""
    except Exception:
        theme = ""
    try:
        brightness = round(float(color_format.brightness), 5)
    except Exception:
        brightness = None
    return (color_type, rgb, theme, brightness)


def run_style_signature(font) -> tuple:
    return (
        font.name,
        round(font.size.pt, 4) if font.size else None,
        bool(font.bold),
        bool(font.italic),
        bool(font.underline),
        int(font.language_id) if font.language_id is not None else None,
        color_signature(font.color),
    )


def text_body_signature(shape) -> tuple:
    text_frame = shape.text_frame
    paragraphs = []
    for paragraph in text_frame.paragraphs:
        runs = tuple((run.text, run_style_signature(run.font)) for run in paragraph.runs)
        paragraphs.append(
            (
                paragraph.level,
                str(paragraph.alignment),
                str(paragraph.line_spacing),
                str(paragraph.space_before),
                str(paragraph.space_after),
                runs,
            )
        )
    return (
        normalize_visible(shape.text),
        (
            text_frame.margin_left,
            text_frame.margin_right,
            text_frame.margin_top,
            text_frame.margin_bottom,
            str(text_frame.vertical_anchor),
            text_frame.word_wrap,
            str(text_frame.auto_size),
        ),
        tuple(paragraphs),
    )


def shape_signature(shape, include_geometry: bool = True) -> tuple:
    text = text_body_signature(shape) if getattr(shape, "has_text_frame", False) else ()
    picture = ""
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        picture = hashlib.sha256(shape.image.blob).hexdigest()
    return (
        str(shape.shape_type),
        placeholder_type(shape),
        rounded_geometry(shape) if include_geometry else (),
        text,
        picture,
    )


def slide_notes(slide) -> tuple[str, ...]:
    try:
        return tuple(
            shape.text
            for shape in slide.notes_slide.shapes
            if getattr(shape, "has_text_frame", False)
        )
    except Exception:
        return ()


def media_hashes(path: Path) -> Counter[str]:
    with zipfile.ZipFile(path) as archive:
        return Counter(
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("ppt/media/")
        )


def package_ok(path: Path) -> tuple[bool, str]:
    try:
        with zipfile.ZipFile(path) as archive:
            broken = archive.testzip()
            names = set(archive.namelist())
            required = {"[Content_Types].xml", "ppt/presentation.xml", "ppt/slides/slide1.xml"}
            missing = sorted(required - names)
            macros = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
        return broken is None and not missing and not macros, (
            f"损坏成员={broken}，缺失={missing}，宏={macros}"
        )
    except Exception as exc:
        return False, f"ZIP 检查失败：{exc}"


def find_roles(presentation: Presentation):
    if len(presentation.slides) != 1:
        return None, None, []
    shapes = list(presentation.slides[0].shapes)
    target_candidates = [
        shape
        for shape in shapes
        if getattr(shape, "has_text_frame", False)
        and normalize_visible(shape.text).startswith("Artificial Intelligence is revolutionizing")
    ]
    title_candidates = [
        shape
        for shape in shapes
        if getattr(shape, "has_text_frame", False) and normalize_visible(shape.text) == "Conclusion"
    ]
    target = target_candidates[0] if len(target_candidates) == 1 else None
    title = title_candidates[0] if len(title_candidates) == 1 else None
    return target, title, shapes


def geometry_checks(source_target, candidate_target, slide_width: int, slide_height: int) -> tuple[bool, dict]:
    left, top, width, height = geometry(candidate_target)
    right, bottom = left + width, top + height
    position_same = (
        abs(left - source_target.left) <= POSITION_TOLERANCE
        and abs(top - source_target.top) <= POSITION_TOLERANCE
    )
    inside = right <= slide_width + POSITION_TOLERANCE and bottom <= slide_height + POSITION_TOLERANCE
    practical_size = width >= int(9.50 * EMU_PER_INCH) and height >= int(3.00 * EMU_PER_INCH)
    resized = width < source_target.width - int(0.50 * EMU_PER_INCH) and height > source_target.height + int(0.30 * EMU_PER_INCH)
    return position_same and inside and practical_size and resized, {
        "source_box_in": [round(value / EMU_PER_INCH, 4) for value in geometry(source_target)],
        "candidate_box_in": [round(value / EMU_PER_INCH, 4) for value in geometry(candidate_target)],
        "position_same": position_same,
        "inside_slide": inside,
        "practical_size": practical_size,
        "width_reduced_height_increased": resized,
    }


def margin_checks(candidate_target, title, slide_width: int, slide_height: int) -> tuple[bool, dict]:
    right_margin = (slide_width - candidate_target.left - candidate_target.width) / EMU_PER_INCH
    bottom_margin = (slide_height - candidate_target.top - candidate_target.height) / EMU_PER_INCH
    title_gap = (candidate_target.top - title.top - title.height) / EMU_PER_INCH
    aspect_ratio = candidate_target.width / candidate_target.height
    passed = (
        0.30 <= right_margin <= 2.60
        and 0.25 <= bottom_margin <= 1.25
        and title_gap >= 0.15
        and 2.7 <= aspect_ratio <= 4.2
    )
    return passed, {
        "right_margin_in": round(right_margin, 4),
        "bottom_margin_in": round(bottom_margin, 4),
        "title_gap_in": round(title_gap, 4),
        "aspect_ratio": round(aspect_ratio, 4),
    }


def render_presentation(path: Path, outdir: Path) -> tuple[Path | None, Path | None, int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "lo-profile"
    profile.mkdir()
    completed = subprocess.run(
        [
            "soffice",
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(outdir),
            str(path),
        ],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return None, None, 0, f"转 PDF 失败：{completed.stdout[-500:]}"
    info = subprocess.run(
        ["pdfinfo", str(pdf)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=30,
        check=False,
    ).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    pages = int(match.group(1)) if match else 0
    image_prefix = outdir / "slide"
    raster = subprocess.run(
        ["pdftoppm", "-singlefile", "-png", "-r", "150", str(pdf), str(image_prefix)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    image = outdir / "slide.png"
    return pdf, image if raster.returncode == 0 and image.is_file() else None, pages, (
        f"页数={pages}，soffice={completed.returncode}，pdftoppm={raster.returncode}"
    )


def extract_pdf_text(pdf: Path) -> str:
    completed = subprocess.run(
        ["pdftotext", "-layout", str(pdf), "-"],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=60,
        check=False,
    )
    return normalize_pdf_text(completed.stdout) if completed.returncode == 0 else ""


def emu_rect_to_pixels(rectangle, image_size, slide_size, padding: int = 0):
    image_width, image_height = image_size
    slide_width, slide_height = slide_size
    left, top, width, height = rectangle
    return (
        max(0, math.floor(left * image_width / slide_width) - padding),
        max(0, math.floor(top * image_height / slide_height) - padding),
        min(image_width, math.ceil((left + width) * image_width / slide_width) + padding),
        min(image_height, math.ceil((top + height) * image_height / slide_height) + padding),
    )


def masked_similarity(first: Path, second: Path, ignored_rectangle, slide_size) -> float:
    with Image.open(first) as first_image, Image.open(second) as second_image:
        left, right = first_image.convert("RGB"), second_image.convert("RGB")
        if left.size != right.size:
            return 0.0
        difference = ImageChops.difference(left, right)
        mask = Image.new("L", left.size, 255)
        draw = ImageDraw.Draw(mask)
        draw.rectangle(emu_rect_to_pixels(ignored_rectangle, left.size, slide_size, padding=10), fill=0)
        stats = ImageStat.Stat(difference, mask=mask)
        rms = statistics.mean(stats.rms) if stats.rms else 255.0
        return max(0.0, 1.0 - rms / 255.0)


def crop_visibility_metrics(image_path: Path, target, slide_size) -> dict:
    with Image.open(image_path) as image:
        gray = image.convert("L")
        crop = gray.crop(emu_rect_to_pixels(geometry(target), gray.size, slide_size))
        histogram = crop.histogram()
        bright_ratio = sum(histogram[180:]) / (crop.width * crop.height)
        stddev = statistics.mean(ImageStat.Stat(crop).stddev)
        return {"bright_ratio": float(bright_ratio), "stddev": float(stddev)}


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {
        "criterion_id": identifier,
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
        "raw": {"description": description, "weight": weight, "type": kind, "passed": bool(passed)},
    }


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [
        item["criterion_id"]
        for item in criteria
        if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]
    ]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(
        item["raw"]["weight"]
        for item in criteria
        if item["raw"]["weight"] > 0 and item["raw"]["passed"]
    )
    penalty = sum(
        abs(item["raw"]["weight"])
        for item in criteria
        if item["raw"]["weight"] < 0 and not item["raw"]["passed"]
    )
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [sha256(path)] if path.is_file() else [],
        "criterion_results": criteria,
        "reward": reward,
        "passed": reward >= 0.8 and not hurdles,
        "hurdle_failures": hurdles,
    }


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        return make_result(path, [result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配") for identifier in SPECS])
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        return make_result(path, [result_item(identifier, False, f"PPTX 打开失败：{exc}") for identifier in SPECS])

    criteria = []
    package_passed, package_evidence = package_ok(path)
    size_passed = (
        len(source.slides) == 1
        and len(candidate.slides) == 1
        and candidate.slide_width == source.slide_width
        and candidate.slide_height == source.slide_height
    )
    criteria.append(
        result_item(
            "R001",
            package_passed and size_passed,
            f"{package_evidence}，输入页数={len(source.slides)}，候选页数={len(candidate.slides)}，"
            f"页面={candidate.slide_width / EMU_PER_INCH:.4f}×{candidate.slide_height / EMU_PER_INCH:.4f}英寸",
        )
    )

    source_target, source_title, source_shapes = find_roles(source)
    target, title, shapes = find_roles(candidate)
    text_passed = (
        source_target is not None
        and target is not None
        and text_body_signature(target) == text_body_signature(source_target)
        and shape_signature(target, include_geometry=False) == shape_signature(source_target, include_geometry=False)
    )
    criteria.append(
        result_item(
            "R002",
            text_passed,
            f"源文本对象={source_target is not None}，候选文本对象={target is not None}，"
            f"内容与文本格式签名一致={text_passed}",
        )
    )

    if source_target is not None and target is not None:
        geometry_passed, geometry_evidence = geometry_checks(
            source_target, target, candidate.slide_width, candidate.slide_height
        )
    else:
        geometry_passed, geometry_evidence = False, {"reason": "缺失源或候选主文本框"}
    criteria.append(result_item("R003", geometry_passed, json.dumps(geometry_evidence, ensure_ascii=False)))

    title_passed = (
        source_title is not None
        and title is not None
        and shape_signature(title) == shape_signature(source_title)
    )
    notes_passed = (
        len(candidate.slides) == 1
        and slide_notes(candidate.slides[0]) == slide_notes(source.slides[0])
    )
    criteria.append(
        result_item(
            "R005",
            title_passed and notes_passed,
            f"标题内容/样式/几何保持={title_passed}，备注保持={notes_passed}",
        )
    )

    full_slide_pictures = []
    if len(candidate.slides) == 1:
        page_area = candidate.slide_width * candidate.slide_height
        full_slide_pictures = [
            shape.name
            for shape in candidate.slides[0].shapes
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
            and shape.width * shape.height / page_area >= 0.85
        ]
    native_passed = (
        target is not None
        and title is not None
        and len(shapes) == 2
        and target.shape_type == MSO_SHAPE_TYPE.TEXT_BOX
        and getattr(title, "has_text_frame", False)
        and not full_slide_pictures
    )
    criteria.append(
        result_item(
            "R006",
            native_passed,
            f"页面对象数={len(shapes)}，主文本类型={str(target.shape_type) if target is not None else '缺失'}，"
            f"标题存在={title is not None}，整页图片={full_slide_pictures}",
        )
    )

    if target is not None and title is not None:
        margin_passed, margin_evidence = margin_checks(
            target, title, candidate.slide_width, candidate.slide_height
        )
    else:
        margin_passed, margin_evidence = False, {"reason": "缺失主文本框或标题"}
    criteria.append(result_item("R007", margin_passed, json.dumps(margin_evidence, ensure_ascii=False)))

    try:
        source_media = media_hashes(source_path)
        candidate_media = media_hashes(path)
        media_passed = all(candidate_media[value] >= count for value, count in source_media.items())
    except Exception:
        source_media, candidate_media, media_passed = Counter(), Counter(), False
    criteria.append(
        result_item(
            "R009",
            package_passed and media_passed,
            f"原媒体保留={media_passed}，媒体数={sum(source_media.values())}→{sum(candidate_media.values())}，{package_evidence}",
        )
    )

    artifact_files = sorted(
        str(item.relative_to(artifact_dir)) for item in artifact_dir.rglob("*") if item.is_file()
    )
    all_text = "\n".join(
        shape.text
        for slide in candidate.slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )
    all_notes = "\n".join(text for slide in candidate.slides for text in slide_notes(slide))
    lowered = f"{all_text}\n{all_notes}".lower()
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker.lower() in lowered]
    hygiene_passed = artifact_files == [OUTPUT_NAME] and not markers
    criteria.append(result_item("R010", hygiene_passed, f"输出文件={artifact_files}，可疑标记={markers}"))

    pre_render_ready = all(
        next(item for item in criteria if item["criterion_id"] == identifier)["raw"]["passed"]
        for identifier in ["R001", "R002", "R003", "R005", "R006", "R009"]
    )
    visibility_passed = False
    visual_passed = False
    extracted_text = ""
    outside_similarity = 0.0
    crop_metrics = {"bright_ratio": 0.0, "stddev": 0.0}
    render_evidence = "前置结构、内容或保持检查失败，未执行渲染"
    if pre_render_ready and source_target is not None and target is not None:
        with tempfile.TemporaryDirectory(prefix="pptarena031-judge-") as temporary:
            root = Path(temporary)
            source_pdf, source_image, source_pages, source_note = render_presentation(source_path, root / "source")
            candidate_pdf, candidate_image, candidate_pages, candidate_note = render_presentation(path, root / "candidate")
            render_evidence = f"源[{source_note}]；候选[{candidate_note}]"
            if candidate_pdf is not None:
                extracted_text = extract_pdf_text(candidate_pdf)
            expected_body = normalize_visible(source_target.text)
            expected_title = normalize_visible(source_title.text) if source_title is not None else "Conclusion"
            visibility_passed = (
                candidate_pages == 1
                and expected_title in extracted_text
                and expected_body in extracted_text
            )
            if source_image is not None and candidate_image is not None:
                ignored = (
                    int(0.50 * EMU_PER_INCH),
                    int(3.40 * EMU_PER_INCH),
                    int(12.83 * EMU_PER_INCH),
                    int(3.75 * EMU_PER_INCH),
                )
                outside_similarity = masked_similarity(
                    source_image,
                    candidate_image,
                    ignored,
                    (candidate.slide_width, candidate.slide_height),
                )
                crop_metrics = crop_visibility_metrics(
                    candidate_image,
                    target,
                    (candidate.slide_width, candidate.slide_height),
                )
            visual_passed = (
                source_pages == 1
                and candidate_pages == 1
                and source_image is not None
                and candidate_image is not None
                and outside_similarity >= 0.9985
                and crop_metrics.get("bright_ratio", 0.0) >= 0.025
                and crop_metrics.get("stddev", 0.0) >= 20.0
                and visibility_passed
            )

    criteria.append(
        result_item(
            "R004",
            visibility_passed,
            f"渲染文本完整包含标题与 813 字符主文={visibility_passed}，"
            f"提取文本长度={len(extracted_text)}；{render_evidence}",
        )
    )
    criteria.append(
        result_item(
            "R008",
            visual_passed,
            f"目标区外像素相似度={outside_similarity:.6f}，"
            f"主文本区亮像素比例={crop_metrics.get('bright_ratio', 0.0):.4f}，"
            f"灰度标准差={crop_metrics.get('stddev', 0.0):.4f}；{render_evidence}",
        )
    )

    severe = (
        not pre_render_ready
        or not visibility_passed
        or not visual_passed
        or bool(full_slide_pictures)
        or bool(markers)
    )
    criteria.append(result_item("R011", not severe, f"严重错误={severe}"))
    order = {identifier: index for index, identifier in enumerate(SPECS)}
    criteria.sort(key=lambda item: order[item["criterion_id"]])
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    output = args.artifact_dir / OUTPUT_NAME
    if output.is_file():
        evaluation = evaluate(output, args.artifact_dir)
    else:
        evaluation = {
            "task_id": TASK_ID,
            "rubric_version": RUBRIC_VERSION,
            "judge_version": JUDGE_VERSION,
            "artifact_sha256": [],
            "criterion_results": [],
            "reward": 0.0,
            "passed": False,
            "hurdle_failures": ["missing_output"],
        }
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
