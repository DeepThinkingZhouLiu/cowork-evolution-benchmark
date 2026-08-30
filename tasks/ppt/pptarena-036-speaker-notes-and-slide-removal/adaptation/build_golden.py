#!/usr/bin/env python3
"""PPTArena Case 36 专用 Golden 构建器。

这个脚本只实现本题已逐页审计的操作：
1. 从原第 2 页按文档顺序取出 8 个非空文本框；
2. 把每个文本框的原始段落 XML 写入原第 3–10 页的备注正文占位符；
3. 从文稿顺序和包中删除原第 2 页及其空备注页。

保留的 slide XML、slide rels、媒体和布局文件均原样复用，避免普通
PowerPoint 库的全文稿重写影响非目标页面。
"""

from __future__ import annotations

import argparse
import hashlib
import posixpath
import re
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation


INPUT_SHA256 = "a81cd08f0bdd05079ba350c1fa32b95db8111e504ae657613707e69abf8fd201"
EXPECTED_TITLES = [
    "Panel to Panel Transitions in Comics",
    "Panel to Panel Transitions",
    "Moment To Moment Transitions",
    "Action to Action Transitions",
    "Subject to Subject Transitions",
    "Scene to scene Transitions",
    "Aspect to Aspect Transitions",
    "Non Sequitur Transitions",
    "Group Activity: Transitions and empathy / Identification",
]

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PR = "http://schemas.openxmlformats.org/package/2006/relationships"
VT = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"
NS = {"p": P, "a": A, "r": R, "pr": PR, "vt": VT}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalize(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def shape_text(shape) -> str:
    return normalize("".join(shape.itertext()))


def read_source_paragraphs(blobs: dict[str, bytes]) -> tuple[list[list[etree._Element]], list[str]]:
    slide = etree.fromstring(blobs["ppt/slides/slide2.xml"])
    groups: list[list[etree._Element]] = []
    texts: list[str] = []
    for shape in slide.xpath(".//p:sp", namespaces=NS):
        text_body = shape.find(f"{{{P}}}txBody")
        if text_body is None or not shape_text(text_body):
            continue
        paragraphs = text_body.findall(f"{{{A}}}p")
        groups.append([deepcopy(paragraph) for paragraph in paragraphs])
        texts.append(shape_text(text_body))
    if len(groups) != 8:
        raise ValueError(f"原第 2 页应有 8 个非空文本框，实际为 {len(groups)}")
    return groups, texts


def remove_slide_two(blobs: dict[str, bytes]) -> set[str]:
    presentation = etree.fromstring(blobs["ppt/presentation.xml"])
    slide_ids = presentation.find(f"{{{P}}}sldIdLst")
    if slide_ids is None or len(slide_ids) != 10:
        raise ValueError("冻结输入不再是 10 页")
    removed = slide_ids[1]
    relation_id = removed.get(f"{{{R}}}id")
    slide_ids.remove(removed)
    blobs["ppt/presentation.xml"] = etree.tostring(
        presentation, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    presentation_rels = etree.fromstring(blobs["ppt/_rels/presentation.xml.rels"])
    matches = [relation for relation in presentation_rels if relation.get("Id") == relation_id]
    if len(matches) != 1 or matches[0].get("Target") != "slides/slide2.xml":
        raise ValueError("原第 2 页关系无法唯一确定")
    presentation_rels.remove(matches[0])
    blobs["ppt/_rels/presentation.xml.rels"] = etree.tostring(
        presentation_rels, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    slide_rels_name = "ppt/slides/_rels/slide2.xml.rels"
    slide_rels = etree.fromstring(blobs[slide_rels_name])
    note_rels = [relation for relation in slide_rels if relation.get("Type", "").endswith("/notesSlide")]
    if len(note_rels) != 1:
        raise ValueError("原第 2 页的备注关系无法唯一确定")
    note_name = posixpath.normpath(posixpath.join("ppt/slides", note_rels[0].get("Target")))
    note_rels_name = f"ppt/notesSlides/_rels/{posixpath.basename(note_name)}.rels"
    return {"ppt/slides/slide2.xml", slide_rels_name, note_name, note_rels_name}


def write_notes(blobs: dict[str, bytes], paragraph_groups: list[list[etree._Element]]) -> None:
    for slide_number, paragraphs in zip(range(3, 11), paragraph_groups):
        rels_name = f"ppt/slides/_rels/slide{slide_number}.xml.rels"
        rels = etree.fromstring(blobs[rels_name])
        note_rels = [relation for relation in rels if relation.get("Type", "").endswith("/notesSlide")]
        if len(note_rels) != 1:
            raise ValueError(f"原第 {slide_number} 页备注关系异常")
        note_name = posixpath.normpath(
            posixpath.join("ppt/slides", note_rels[0].get("Target"))
        )
        notes = etree.fromstring(blobs[note_name])
        body_shapes = notes.xpath(
            './/p:sp[p:nvSpPr/p:nvPr/p:ph[@type="body"]]', namespaces=NS
        )
        if len(body_shapes) != 1:
            raise ValueError(f"{note_name} 备注正文占位符异常")
        text_body = body_shapes[0].find(f"{{{P}}}txBody")
        for old in list(text_body.findall(f"{{{A}}}p")):
            text_body.remove(old)
        for paragraph in paragraphs:
            text_body.append(deepcopy(paragraph))
        blobs[note_name] = etree.tostring(
            notes, xml_declaration=True, encoding="UTF-8", standalone=True
        )


def update_content_types(blobs: dict[str, bytes], removed_names: set[str]) -> None:
    content_types = etree.fromstring(blobs["[Content_Types].xml"])
    for override in list(content_types):
        if override.get("PartName", "").lstrip("/") in removed_names:
            content_types.remove(override)
    blobs["[Content_Types].xml"] = etree.tostring(
        content_types, xml_declaration=True, encoding="UTF-8", standalone=True
    )


def update_app_metadata(blobs: dict[str, bytes]) -> None:
    properties = etree.fromstring(blobs["docProps/app.xml"])
    for element in properties.iter():
        local_name = etree.QName(element).localname
        if local_name in {"Slides", "Notes"}:
            element.text = "9"

    heading_vectors = properties.xpath(
        './/*[local-name()="HeadingPairs"]/vt:vector', namespaces=NS
    )
    if len(heading_vectors) == 1:
        variants = heading_vectors[0].findall(f"{{{VT}}}variant")
        for index, variant in enumerate(variants[:-1]):
            labels = variant.findall(f"{{{VT}}}lpstr")
            if labels and normalize(labels[0].text or "") == "Slide Titles":
                counts = variants[index + 1].findall(f"{{{VT}}}i4")
                if len(counts) == 1:
                    counts[0].text = "9"

    vectors = properties.xpath('.//*[local-name()="TitlesOfParts"]/vt:vector', namespaces=NS)
    if len(vectors) == 1:
        values = vectors[0].findall(f"{{{VT}}}lpstr")
        matches = [value for value in values if normalize(value.text or "") == "PowerPoint Presentation"]
        if len(matches) == 1:
            vectors[0].remove(matches[0])
            vectors[0].set("size", str(len(values) - 1))

    blobs["docProps/app.xml"] = etree.tostring(
        properties, xml_declaration=True, encoding="UTF-8", standalone=True
    )


def validate_output(path: Path, expected_notes: list[str]) -> None:
    presentation = Presentation(path)
    if len(presentation.slides) != 9:
        raise ValueError("产物页数不是 9")
    actual_titles = [normalize(slide.shapes.title.text if slide.shapes.title else "") for slide in presentation.slides]
    if actual_titles != EXPECTED_TITLES:
        raise ValueError(f"产物顺序或标题异常：{actual_titles}")
    actual_notes = [normalize(slide.notes_slide.notes_text_frame.text) for slide in presentation.slides]
    if actual_notes[0]:
        raise ValueError("封面不应有讲者备注")
    if actual_notes[1:] != expected_notes:
        raise ValueError("备注文字未按原顺序完整转移")


def build(source: Path, output: Path) -> None:
    if sha256(source) != INPUT_SHA256:
        raise ValueError("冻结输入 SHA256 不匹配")
    source_presentation = Presentation(source)
    expected_notes = [
        normalize(shape.text)
        for shape in source_presentation.slides[1].shapes
        if getattr(shape, "has_text_frame", False) and normalize(shape.text)
    ]
    if len(expected_notes) != 8:
        raise ValueError("无法从冻结输入读取 8 段备注正文")
    with ZipFile(source) as archive:
        infos = {info.filename: info for info in archive.infolist()}
        blobs = {name: archive.read(name) for name in infos}

    paragraphs, _ = read_source_paragraphs(blobs)
    write_notes(blobs, paragraphs)
    removed_names = remove_slide_two(blobs)
    update_content_types(blobs, removed_names)
    update_app_metadata(blobs)

    output.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(output, "w") as archive:
        for name, info in infos.items():
            if name in removed_names:
                continue
            archive.writestr(info, blobs[name])
    validate_output(output, expected_notes)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.source, args.output)
    print(f"{args.output}\nsha256={sha256(args.output)}")


if __name__ == "__main__":
    main()
