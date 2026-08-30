#!/usr/bin/env python3
"""只为 PPTArena Case 36 构造合法变体、负例和中等例。"""

from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches


OUTPUT_NAME = "AddSpeakerNotes_GroundTruth.pptx"


def output_path(root: Path, case: str) -> Path:
    target = root / case / OUTPUT_NAME
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def save_copy(golden: Path, output: Path, mutator) -> None:
    presentation = Presentation(golden)
    mutator(presentation)
    presentation.save(output)


def valid_alternate(golden: Path, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        frame = presentation.slides[2].notes_slide.notes_text_frame
        frame.text = frame.text.replace(". ", ".\n", 1)

    save_copy(golden, output, mutate)


def extra_note(golden: Path, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        frame = presentation.slides[1].notes_slide.notes_text_frame
        frame.text = f"{frame.text}\nPresenter cue: pause before the next example."

    save_copy(golden, output, mutate)


def missing_note(golden: Path, output: Path) -> None:
    save_copy(golden, output, lambda presentation: setattr(
        presentation.slides[4].notes_slide.notes_text_frame, "text", ""
    ))


def swapped_notes(golden: Path, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        first = presentation.slides[2].notes_slide.notes_text_frame
        second = presentation.slides[3].notes_slide.notes_text_frame
        first_text, second_text = first.text, second.text
        first.text, second.text = second_text, first_text

    save_copy(golden, output, mutate)


def truncated_note(golden: Path, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        frame = presentation.slides[6].notes_slide.notes_text_frame
        frame.text = frame.text[: len(frame.text) // 2]

    save_copy(golden, output, mutate)


def duplicated_note(golden: Path, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        source = presentation.slides[1].notes_slide.notes_text_frame.text
        target = presentation.slides[2].notes_slide.notes_text_frame
        target.text = f"{target.text}\n{source}"

    save_copy(golden, output, mutate)


def title_note(golden: Path, output: Path) -> None:
    save_copy(golden, output, lambda presentation: setattr(
        presentation.slides[0].notes_slide.notes_text_frame,
        "text",
        "This title slide should not have speaker notes.",
    ))


def changed_slide_text(golden: Path, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        presentation.slides[2].shapes.title.text = "Moment Transitions"

    save_copy(golden, output, mutate)


def moved_shape(golden: Path, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        presentation.slides[3].shapes[1].left += Inches(0.25)

    save_copy(golden, output, mutate)


def reordered_slides(golden: Path, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        slide_ids = presentation.slides._sldIdLst
        ordered = list(slide_ids)
        ordered[2], ordered[3] = ordered[3], ordered[2]
        for slide_id in list(slide_ids):
            slide_ids.remove(slide_id)
        for slide_id in ordered:
            slide_ids.append(slide_id)

    save_copy(golden, output, mutate)


def added_shape(golden: Path, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        textbox = presentation.slides[1].shapes.add_textbox(
            Inches(7.5), Inches(5.1), Inches(2.0), Inches(0.35)
        )
        textbox.text = "extra"

    save_copy(golden, output, mutate)


def flattened_slide(golden: Path, slide_image: Path, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        slide = presentation.slides[2]
        for shape in list(slide.shapes):
            shape._element.getparent().remove(shape._element)
        slide.shapes.add_picture(
            str(slide_image),
            0,
            0,
            width=presentation.slide_width,
            height=presentation.slide_height,
        )

    save_copy(golden, output, mutate)


def removed_picture(golden: Path, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        for slide in presentation.slides:
            pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
            if pictures:
                pictures[0]._element.getparent().remove(pictures[0]._element)
                return
        raise ValueError("无法定位图片负例")

    save_copy(golden, output, mutate)


def hardcoded_marker(golden: Path, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        frame = presentation.slides[1].notes_slide.notes_text_frame
        frame.text = f"{frame.text}\nsolution/Golden"

    save_copy(golden, output, mutate)


def extra_slide(golden: Path, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        presentation.slides.add_slide(presentation.slide_layouts[0])

    save_copy(golden, output, mutate)


def wrong_slide_deleted(source: Path, output: Path) -> None:
    presentation = Presentation(source)
    slide_id = presentation.slides._sldIdLst[2]
    presentation.part.drop_rel(slide_id.rId)
    del presentation.slides._sldIdLst[2]
    presentation.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--slide-image", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()

    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))
    valid_alternate(args.golden, output_path(args.outdir, "valid-alternate"))
    extra_note(args.golden, output_path(args.outdir, "extra-note"))
    missing_note(args.golden, output_path(args.outdir, "missing-note"))
    swapped_notes(args.golden, output_path(args.outdir, "swapped-notes"))
    truncated_note(args.golden, output_path(args.outdir, "truncated-note"))
    duplicated_note(args.golden, output_path(args.outdir, "duplicated-note"))
    title_note(args.golden, output_path(args.outdir, "title-note"))
    changed_slide_text(args.golden, output_path(args.outdir, "changed-slide-text"))
    moved_shape(args.golden, output_path(args.outdir, "moved-shape"))
    reordered_slides(args.golden, output_path(args.outdir, "reordered-slides"))
    added_shape(args.golden, output_path(args.outdir, "added-shape"))
    flattened_slide(args.golden, args.slide_image, output_path(args.outdir, "flattened-slide"))
    removed_picture(args.golden, output_path(args.outdir, "removed-picture"))
    hardcoded_marker(args.golden, output_path(args.outdir, "hardcoded-marker"))
    extra_slide(args.golden, output_path(args.outdir, "extra-slide"))
    wrong_slide_deleted(args.source, output_path(args.outdir, "wrong-slide-deleted"))


if __name__ == "__main__":
    main()
