#!/usr/bin/env python3
"""PPTArena Case 36 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import posixpath
import re
import statistics
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from PIL import Image, ImageStat
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "AddSpeakerNotes_GroundTruth.pptx"
INPUT_NAME = "AddSpeakerNotes_TestA.pptx"
INPUT_SHA256 = "a81cd08f0bdd05079ba350c1fa32b95db8111e504ae657613707e69abf8fd201"
TASK_ID = "cowork-evo/pptarena-036-speaker-notes-and-slide-removal"
RUBRIC_VERSION = "pptarena-case-036-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-036-deterministic-v0.1"
SUSPICIOUS_MARKERS = [
    "solution/golden",
    "judge-result",
    "artifact_dir",
    "pptarena-case-036",
    "/root/",
]

EXPECTED_TITLES = [
    "Panel to Panel Transitions in Comics",
    "Panel to Panel Transitions",
    "Moment To Moment Transitions",
    "Action to Action Transitions",
    "Subject to Subject Transitions",
    "Scene to scene Transitions",
    "Aspect to Aspect Transitions",
    "Non Sequitur Transitions",
    "Group Activity: Transitions and empathy / Identification",
]

SPECS = {
    "R001": ("PPTX 完整、尺寸不变且严格为 9 页", 7, "hurdle"),
    "R002": ("原第 2 页已删除且保留页顺序正确", 13, "hurdle"),
    "R003": ("8 段原文完整、唯一地映射到对应讲者备注", 24, "hurdle"),
    "R004": ("封面无备注且临时文本框不再可见", 8, "hurdle"),
    "R005": ("保留页的内容、对象、几何、文本样式和页面尺寸一致", 16, "hurdle"),
    "R006": ("原媒体和原生可编辑对象保留，无整页截图", 8, "hurdle"),
    "R007": ("备注无额外断言、提示或夹带内容", 6, "positive"),
    "R008": ("LibreOffice 可渲染 9 页且页面可见内容完整", 8, "hurdle"),
    "R009": ("PPTX 关系完整，9 个幻灯片关系和 9 个备注页有效", 6, "hurdle"),
    "R010": ("输出目录干净且无答案标记", 4, "positive"),
    "R011": ("无删错页、串页、改动原页、丢媒体或丧失可编辑性等严重错误", -20, "penalty"),
}


def normalize(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def locate_source() -> Path | None:
    candidates = [
        Path("/app/input_files") / INPUT_NAME,
        Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME,
    ]
    return next((path for path in candidates if path.is_file()), None)


def safe_value(callback, default=""):
    try:
        return callback()
    except Exception:
        return default


def color_signature(color_format) -> tuple:
    return tuple(
        safe_value(lambda name=name: str(getattr(color_format, name)))
        for name in ["type", "rgb", "theme_color", "brightness"]
    )


def font_signature(font) -> tuple:
    return (
        font.name,
        round(font.size.pt, 4) if font.size else None,
        font.bold,
        font.italic,
        font.underline,
        safe_value(lambda: str(font.language_id)),
        color_signature(font.color),
    )


def text_signature(shape) -> tuple:
    if not getattr(shape, "has_text_frame", False):
        return ()
    frame = shape.text_frame
    paragraphs = []
    for paragraph in frame.paragraphs:
        paragraphs.append(
            (
                paragraph.level,
                str(paragraph.alignment),
                str(paragraph.line_spacing),
                str(paragraph.space_before),
                str(paragraph.space_after),
                tuple((run.text, font_signature(run.font)) for run in paragraph.runs),
            )
        )
    return (
        normalize(shape.text),
        (
            frame.margin_left,
            frame.margin_right,
            frame.margin_top,
            frame.margin_bottom,
            str(frame.vertical_anchor),
            frame.word_wrap,
            str(frame.auto_size),
        ),
        tuple(paragraphs),
    )


def shape_signature(shape) -> tuple:
    picture_hash = ""
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        picture_hash = hashlib.sha256(shape.image.blob).hexdigest()
    placeholder = str(shape.placeholder_format.type) if shape.is_placeholder else ""
    geometry = tuple(round(int(value) / 1000) * 1000 for value in (shape.left, shape.top, shape.width, shape.height))
    return (
        str(shape.shape_type),
        placeholder,
        geometry,
        safe_value(lambda: round(float(shape.rotation), 4), 0.0),
        text_signature(shape),
        picture_hash,
    )


def slide_signature(slide) -> tuple:
    return tuple(shape_signature(shape) for shape in slide.shapes)


def notes_body(slide) -> str:
    try:
        return normalize(slide.notes_slide.notes_text_frame.text)
    except Exception:
        return ""


def visible_text(slide) -> str:
    return normalize(
        "\n".join(
            shape.text
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and shape.text
        )
    )


def media_hashes(path: Path) -> Counter[str]:
    with zipfile.ZipFile(path) as archive:
        return Counter(
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("ppt/media/")
        )


def relationship_owner(rels_name: str) -> str:
    if rels_name == "_rels/.rels":
        return ""
    marker = "/_rels/"
    prefix, filename = rels_name.split(marker, 1)
    return posixpath.join(prefix, filename[:-5])


def package_checks(path: Path) -> tuple[bool, dict]:
    evidence = {
        "broken_member": None,
        "missing_targets": [],
        "macros": [],
        "slide_relations": 0,
        "notes_parts": 0,
    }
    try:
        with zipfile.ZipFile(path) as archive:
            evidence["broken_member"] = archive.testzip()
            names = set(archive.namelist())
            required = {"[Content_Types].xml", "ppt/presentation.xml", "ppt/_rels/presentation.xml.rels"}
            evidence["required_missing"] = sorted(required - names)
            evidence["macros"] = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
            evidence["notes_parts"] = len(
                [name for name in names if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", name)]
            )
            for rels_name in sorted(name for name in names if name.endswith(".rels")):
                try:
                    root = __import__("xml.etree.ElementTree", fromlist=["ElementTree"]).fromstring(
                        archive.read(rels_name)
                    )
                except Exception as exc:
                    evidence["missing_targets"].append(f"{rels_name}: XML 错误 {exc}")
                    continue
                owner = relationship_owner(rels_name)
                base = posixpath.dirname(owner)
                for relation in root:
                    relation_type = relation.attrib.get("Type", "")
                    if rels_name == "ppt/_rels/presentation.xml.rels" and relation_type.endswith("/slide"):
                        evidence["slide_relations"] += 1
                    if relation.attrib.get("TargetMode") == "External":
                        continue
                    target = relation.attrib.get("Target", "")
                    resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                    if resolved not in names:
                        evidence["missing_targets"].append(f"{rels_name} -> {resolved}")
        passed = (
            evidence["broken_member"] is None
            and not evidence["required_missing"]
            and not evidence["macros"]
            and not evidence["missing_targets"]
        )
        return passed, evidence
    except Exception as exc:
        evidence["exception"] = str(exc)
        return False, evidence


def render_checks(path: Path, expected_titles: list[str]) -> tuple[bool, dict]:
    evidence = {"pages": 0, "images": 0, "nonwhite_ratios": [], "titles_found": []}
    try:
        with tempfile.TemporaryDirectory(prefix="pptarena036-judge-") as temporary:
            root = Path(temporary)
            profile = root / "lo-profile"
            profile.mkdir()
            completed = subprocess.run(
                [
                    "soffice",
                    f"-env:UserInstallation=file://{profile}",
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(root),
                    str(path),
                ],
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                timeout=300,
                check=False,
            )
            pdf = root / f"{path.stem}.pdf"
            if completed.returncode != 0 or not pdf.is_file():
                evidence["error"] = f"PDF 转换失败：{completed.stdout[-500:]}"
                return False, evidence
            info = subprocess.run(
                ["pdfinfo", str(pdf)],
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                timeout=60,
                check=False,
            ).stdout
            match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
            evidence["pages"] = int(match.group(1)) if match else 0
            text_result = subprocess.run(
                ["pdftotext", "-layout", str(pdf), "-"],
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                timeout=90,
                check=False,
            )
            pdf_text = normalize(text_result.stdout) if text_result.returncode == 0 else ""
            evidence["titles_found"] = [title for title in expected_titles if normalize(title) in pdf_text]
            raster = subprocess.run(
                ["pdftoppm", "-png", "-r", "96", str(pdf), str(root / "slide")],
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                timeout=300,
                check=False,
            )
            images = sorted(root.glob("slide-*.png")) if raster.returncode == 0 else []
            evidence["images"] = len(images)
            for image_path in images:
                with Image.open(image_path) as image:
                    gray = image.convert("L")
                    histogram = gray.histogram()
                    nonwhite = sum(histogram[:245]) / max(1, gray.width * gray.height)
                    evidence["nonwhite_ratios"].append(round(nonwhite, 6))
            passed = (
                evidence["pages"] == 9
                and evidence["images"] == 9
                and len(evidence["titles_found"]) == 9
                and all(value >= 0.003 for value in evidence["nonwhite_ratios"])
            )
            return passed, evidence
    except Exception as exc:
        evidence["exception"] = str(exc)
        return False, evidence


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {
        "criterion_id": identifier,
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
        "raw": {"description": description, "weight": weight, "type": kind, "passed": bool(passed)},
    }


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [
        item["criterion_id"]
        for item in criteria
        if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]
    ]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(
        item["raw"]["weight"]
        for item in criteria
        if item["raw"]["weight"] > 0 and item["raw"]["passed"]
    )
    penalty = sum(
        abs(item["raw"]["weight"])
        for item in criteria
        if item["raw"]["weight"] < 0 and not item["raw"]["passed"]
    )
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [sha256(path)] if path.is_file() else [],
        "criterion_results": criteria,
        "reward": reward,
        "passed": reward >= 0.8 and not hurdles,
        "hurdle_failures": hurdles,
    }


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        return make_result(path, [result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配") for identifier in SPECS])
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        return make_result(path, [result_item(identifier, False, f"PPTX 打开失败：{exc}") for identifier in SPECS])

    criteria: list[dict] = []
    package_passed, package_evidence = package_checks(path)
    size_passed = (
        len(source.slides) == 10
        and len(candidate.slides) == 9
        and candidate.slide_width == source.slide_width
        and candidate.slide_height == source.slide_height
    )
    criteria.append(
        result_item(
            "R001",
            package_passed and size_passed,
            f"输入页数={len(source.slides)}，候选页数={len(candidate.slides)}，"
            f"页面尺寸一致={candidate.slide_width == source.slide_width and candidate.slide_height == source.slide_height}，"
            f"包检查={json.dumps(package_evidence, ensure_ascii=False)}",
        )
    )

    actual_titles = [normalize(slide.shapes.title.text if slide.shapes.title else "") for slide in candidate.slides]
    order_passed = actual_titles == EXPECTED_TITLES
    criteria.append(result_item("R002", order_passed, f"实际标题顺序={actual_titles}"))

    expected_notes = [
        normalize(shape.text)
        for shape in source.slides[1].shapes
        if getattr(shape, "has_text_frame", False) and normalize(shape.text)
    ]
    actual_notes = [notes_body(slide) for slide in candidate.slides]
    mapping_details = []
    mapping_passed = len(expected_notes) == 8 and len(actual_notes) == 9
    if mapping_passed:
        for index, expected in enumerate(expected_notes, start=1):
            matching_pages = [page for page, note in enumerate(actual_notes) if expected in note]
            page_passed = expected in actual_notes[index] and matching_pages == [index]
            mapping_details.append(
                {
                    "destination_slide": index + 1,
                    "expected_chars": len(expected),
                    "actual_chars": len(actual_notes[index]),
                    "matching_zero_based_pages": matching_pages,
                    "passed": page_passed,
                }
            )
            mapping_passed = mapping_passed and page_passed
    criteria.append(result_item("R003", mapping_passed, json.dumps(mapping_details, ensure_ascii=False)))

    all_visible = [visible_text(slide) for slide in candidate.slides]
    temporary_text_absent = all(
        expected[:80] not in page_text
        for expected in expected_notes
        for page_text in all_visible
    )
    title_note_empty = bool(actual_notes) and not actual_notes[0]
    criteria.append(
        result_item(
            "R004",
            title_note_empty and temporary_text_absent,
            f"封面备注为空={title_note_empty}，临时页长文不在可见页面={temporary_text_absent}",
        )
    )

    source_indices = [0] + list(range(2, 10))
    preservation_details = []
    preservation_passed = len(candidate.slides) == 9
    if preservation_passed:
        for candidate_index, source_index in enumerate(source_indices):
            source_signature = slide_signature(source.slides[source_index])
            candidate_signature = slide_signature(candidate.slides[candidate_index])
            passed = source_signature == candidate_signature
            preservation_details.append(
                {
                    "candidate_slide": candidate_index + 1,
                    "source_slide": source_index + 1,
                    "source_shapes": len(source_signature),
                    "candidate_shapes": len(candidate_signature),
                    "passed": passed,
                }
            )
            preservation_passed = preservation_passed and passed
    criteria.append(result_item("R005", preservation_passed, json.dumps(preservation_details, ensure_ascii=False)))

    source_media = media_hashes(source_path)
    candidate_media = media_hashes(path)
    full_slide_pictures = []
    if len(candidate.slides) == 9:
        page_area = candidate.slide_width * candidate.slide_height
        for page_number, slide in enumerate(candidate.slides, start=1):
            for shape in slide.shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.width * shape.height / page_area >= 0.85:
                    full_slide_pictures.append(f"{page_number}:{shape.name}")
    editable_passed = source_media == candidate_media and not full_slide_pictures and preservation_passed
    criteria.append(
        result_item(
            "R006",
            editable_passed,
            f"媒体哈希集合一致={source_media == candidate_media}，媒体数={sum(source_media.values())}→{sum(candidate_media.values())}，"
            f"整页图片={full_slide_pictures}",
        )
    )

    clean_notes = len(expected_notes) == 8 and len(actual_notes) == 9 and not actual_notes[0]
    if clean_notes:
        clean_notes = all(actual_notes[index] == expected for index, expected in enumerate(expected_notes, start=1))
    criteria.append(result_item("R007", clean_notes, f"备注精确文本对应={clean_notes}"))

    render_ready = all(
        next(item for item in criteria if item["criterion_id"] == identifier)["raw"]["passed"]
        for identifier in ["R001", "R002", "R003", "R004", "R005", "R006"]
    )
    if render_ready:
        render_passed, render_evidence = render_checks(path, EXPECTED_TITLES)
    else:
        render_passed, render_evidence = False, {"skipped": "前置结构或保持检查失败"}
    criteria.append(result_item("R008", render_passed, json.dumps(render_evidence, ensure_ascii=False)))

    relations_passed = (
        package_passed
        and package_evidence.get("slide_relations") == 9
        and package_evidence.get("notes_parts") == 9
    )
    criteria.append(result_item("R009", relations_passed, json.dumps(package_evidence, ensure_ascii=False)))

    artifact_files = sorted(
        str(item.relative_to(artifact_dir)) for item in artifact_dir.rglob("*") if item.is_file()
    )
    searchable = "\n".join(all_visible + actual_notes).lower()
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker.lower() in searchable]
    hygiene_passed = artifact_files == [OUTPUT_NAME] and not markers
    criteria.append(result_item("R010", hygiene_passed, f"输出文件={artifact_files}，可疑标记={markers}"))

    severe = not all(
        next(item for item in criteria if item["criterion_id"] == identifier)["raw"]["passed"]
        for identifier in ["R001", "R002", "R003", "R004", "R005", "R006", "R008", "R009"]
    ) or bool(markers)
    criteria.append(result_item("R011", not severe, f"严重错误={severe}"))
    order = {identifier: index for index, identifier in enumerate(SPECS)}
    criteria.sort(key=lambda item: order[item["criterion_id"]])
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    output = args.artifact_dir / OUTPUT_NAME
    if output.is_file():
        evaluation = evaluate(output, args.artifact_dir)
    else:
        evaluation = {
            "task_id": TASK_ID,
            "rubric_version": RUBRIC_VERSION,
            "judge_version": JUDGE_VERSION,
            "artifact_sha256": [],
            "criterion_results": [],
            "reward": 0.0,
            "passed": False,
            "hurdle_failures": ["missing_output"],
        }
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
