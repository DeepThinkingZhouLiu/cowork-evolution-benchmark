#!/usr/bin/env python3
"""只为 PPTArena Case 37 构造可重放的合法变体、负例和中等例。"""

from __future__ import annotations

import argparse
import shutil
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches


OUTPUT_NAME = "AddTransitions_GroundTruthA.pptx"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def qn(namespace: str, local: str) -> str:
    return f"{{{namespace}}}{local}"


def local_name(element) -> str:
    return etree.QName(element).localname


def output_path(root: Path, case: str) -> Path:
    target = root / case / OUTPUT_NAME
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def slide_number(name: str) -> int | None:
    if not name.startswith("ppt/slides/slide") or not name.endswith(".xml"):
        return None
    value = name.removeprefix("ppt/slides/slide").removesuffix(".xml")
    return int(value) if value.isdigit() else None


def serialize_xml(root) -> bytes:
    return etree.tostring(
        root,
        encoding="UTF-8",
        xml_declaration=True,
        standalone=True,
    )


def replace_transition(
    root,
    effect: str | None,
    *,
    speed: str | None = None,
    through_black: bool = False,
    automatic_ms: int | None = None,
    click_enabled: bool = True,
    sound: bool = False,
    duplicate: bool = False,
) -> None:
    for transition in root.xpath("./p:transition", namespaces={"p": P_NS}):
        transition.getparent().remove(transition)
    if effect is None:
        return

    def make_transition():
        transition = etree.Element(qn(P_NS, "transition"))
        if speed is not None:
            transition.set("spd", speed)
        if automatic_ms is not None:
            transition.set("advTm", str(automatic_ms))
        if not click_enabled:
            transition.set("advClick", "0")
        child = etree.SubElement(transition, qn(P_NS, effect))
        if effect == "fade" and through_black:
            child.set("thruBlk", "1")
        if sound:
            etree.SubElement(transition, qn(P_NS, "sndAc"))
        return transition

    insertion = len(root)
    for index, child in enumerate(root):
        if local_name(child.tag) in {"timing", "extLst"}:
            insertion = index
            break
    root.insert(insertion, make_transition())
    if duplicate:
        root.insert(insertion + 1, make_transition())


def rewrite_archive(source: Path, output: Path, transform) -> None:
    with zipfile.ZipFile(source) as archive_in, zipfile.ZipFile(output, "w") as archive_out:
        for info in archive_in.infolist():
            data = archive_in.read(info.filename)
            data = transform(info.filename, data)
            archive_out.writestr(info, data)


def all_fade(
    source: Path,
    output: Path,
    *,
    speed: str | None = None,
    through_black: bool = False,
) -> None:
    def transform(name: str, data: bytes) -> bytes:
        number = slide_number(name)
        if number is None:
            return data
        root = etree.fromstring(data)
        replace_transition(
            root,
            "fade",
            speed=speed,
            through_black=through_black,
        )
        return serialize_xml(root)

    rewrite_archive(source, output, transform)


def selective_transition(source: Path, output: Path, mode: str) -> None:
    def transform(name: str, data: bytes) -> bytes:
        number = slide_number(name)
        if number is None:
            return data
        root = etree.fromstring(data)
        kwargs = {"effect": "fade", "speed": "slow"}
        if mode == "missing-one" and number == 9:
            kwargs["effect"] = None
        elif mode == "one-cut" and number == 4:
            kwargs["effect"] = "cut"
        elif mode == "all-cut":
            kwargs["effect"] = "cut"
        elif mode == "duplicate" and number == 5:
            kwargs["duplicate"] = True
        elif mode == "auto-advance" and number == 3:
            kwargs["automatic_ms"] = 5000
        elif mode == "disable-click" and number == 3:
            kwargs["click_enabled"] = False
        elif mode == "sound" and number == 3:
            kwargs["sound"] = True
        replace_transition(root, **kwargs)
        return serialize_xml(root)

    rewrite_archive(source, output, transform)


def changed_text(source: Path, output: Path) -> None:
    def transform(name: str, data: bytes) -> bytes:
        number = slide_number(name)
        if number is None:
            return data
        root = etree.fromstring(data)
        replace_transition(root, "fade", speed="slow")
        if number == 2:
            nodes = root.xpath(".//a:t", namespaces={"a": A_NS})
            if not nodes:
                raise ValueError("第 2 页无法定位文字")
            nodes[0].text = f"{nodes[0].text} X"
        return serialize_xml(root)

    rewrite_archive(source, output, transform)


def changed_notes(golden: Path, output: Path) -> None:
    changed = False

    def transform(name: str, data: bytes) -> bytes:
        nonlocal changed
        if name != "ppt/notesSlides/notesSlide3.xml":
            return data
        root = etree.fromstring(data)
        nodes = [node for node in root.xpath(".//a:t", namespaces={"a": A_NS}) if node.text]
        if not nodes:
            raise ValueError("第 3 页备注无可变异文字")
        nodes[-1].text = f"{nodes[-1].text} X"
        changed = True
        return serialize_xml(root)

    rewrite_archive(golden, output, transform)
    if not changed:
        raise ValueError("未找到 notesSlide3.xml")


def added_animation(golden: Path, output: Path) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if slide_number(name) != 3:
            return data
        root = etree.fromstring(data)
        etree.SubElement(root, qn(P_NS, "timing"))
        return serialize_xml(root)

    rewrite_archive(golden, output, transform)


def changed_image(golden: Path, output: Path) -> None:
    changed = False

    def transform(name: str, data: bytes) -> bytes:
        nonlocal changed
        if not changed and name.startswith("ppt/media/"):
            changed = True
            return b"not-the-original-image\n"
        return data

    rewrite_archive(golden, output, transform)
    if not changed:
        raise ValueError("文稿中没有可变异媒体")


def repacked_content_types(golden: Path, output: Path) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name == "[Content_Types].xml":
            return data.replace(b"</Types>", b"  \n</Types>")
        return data

    rewrite_archive(golden, output, transform)


def flattened_slide(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    source_picture = next(
        shape
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    )
    suffix = f".{source_picture.image.ext}"
    with tempfile.NamedTemporaryFile(suffix=suffix) as temporary:
        temporary.write(source_picture.image.blob)
        temporary.flush()
        slide = presentation.slides[0]
        for shape in list(slide.shapes):
            shape._element.getparent().remove(shape._element)
        slide.shapes.add_picture(
            temporary.name,
            0,
            0,
            width=presentation.slide_width,
            height=presentation.slide_height,
        )
        presentation.save(output)


def hardcoded_marker(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    box = presentation.slides[0].shapes.add_textbox(
        Inches(0.10), Inches(5.35), Inches(2.50), Inches(0.20)
    )
    box.text = "solution/Golden"
    presentation.save(output)


def extra_slide(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    presentation.slides.add_slide(presentation.slide_layouts[0])
    presentation.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()

    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))
    all_fade(args.source, output_path(args.outdir, "valid-default"))
    all_fade(
        args.source,
        output_path(args.outdir, "valid-through-black-fast"),
        speed="fast",
        through_black=True,
    )
    repacked_content_types(
        args.golden,
        output_path(args.outdir, "repacked-content-types"),
    )
    for mode in [
        "missing-one",
        "one-cut",
        "all-cut",
        "duplicate",
        "auto-advance",
        "disable-click",
        "sound",
    ]:
        selective_transition(args.source, output_path(args.outdir, mode), mode)
    changed_text(args.source, output_path(args.outdir, "changed-text"))
    changed_notes(args.golden, output_path(args.outdir, "changed-notes"))
    added_animation(args.golden, output_path(args.outdir, "added-animation"))
    changed_image(args.golden, output_path(args.outdir, "changed-image"))
    flattened_slide(args.golden, output_path(args.outdir, "flattened-slide"))
    hardcoded_marker(args.golden, output_path(args.outdir, "hardcoded-marker"))
    extra_slide(args.golden, output_path(args.outdir, "extra-slide"))


if __name__ == "__main__":
    main()
