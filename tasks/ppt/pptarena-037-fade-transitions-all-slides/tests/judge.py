#!/usr/bin/env python3
"""PPTArena Case 37 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from lxml import etree
from PIL import Image, ImageChops
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "AddTransitions_GroundTruthA.pptx"
INPUT_NAME = "AddTransitions_TestA.pptx"
INPUT_SHA256 = "a5557a252a8b8668a916a106b54c9935dd5d1eef2eb76d5f2fa4efa5bd4529b4"
TASK_ID = "cowork-evo/pptarena-037-fade-transitions-all-slides"
RUBRIC_VERSION = "pptarena-case-037-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-037-deterministic-v0.1"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
EMU_PER_INCH = 914_400
SUSPICIOUS_MARKERS = [
    "solution/golden",
    "/root/",
    "judge-result",
    "artifact_dir",
    "pptarena-case-037",
]

SPECS = {
    "R001": ("指定 PPTX 可打开，页面尺寸不变且总页数为 9", 6, "hurdle"),
    "R002": ("全部 9 页都在幻灯片级别恰有一个 Fade 转场", 30, "hurdle"),
    "R003": ("Fade 只使用合理变体，无自动换页、声音或额外效果", 10, "hurdle"),
    "R004": ("全部页面内容、样式、对象、图片、几何、顺序和版式保持", 16, "hurdle"),
    "R005": ("讲者备注和原有动画/时序完整保留", 10, "hurdle"),
    "R006": ("原媒体、页面对象及版式关系保留，无整页截图", 8, "hurdle"),
    "R007": ("9 页可渲染且静态视觉与冻结输入保持", 8, "hurdle"),
    "R008": ("文稿原生可编辑、无宏和损坏关系，转场位于幻灯片 OOXML", 5, "hurdle"),
    "R009": ("除转场和文档元数据外，内容 XML 和关键包部件字节级保持", 5, "positive"),
    "R010": ("输出目录干净且不含答案路径或未授权文字", 2, "positive"),
    "R011": ("任何页缺少 Fade、转场错误、原内容变化或可编辑性破坏", -20, "penalty"),
}


def normalize_visible(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def locate_source() -> Path | None:
    candidates = [
        Path("/app/input_files") / INPUT_NAME,
        Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME,
    ]
    return next((path for path in candidates if path.is_file()), None)


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def placeholder_type(shape) -> str:
    if not getattr(shape, "is_placeholder", False):
        return ""
    return str(shape.placeholder_format.type)


def rounded_geometry(shape) -> tuple:
    return tuple(
        round(int(value) / 1000) * 1000
        for value in (shape.left, shape.top, shape.width, shape.height)
    )


def color_signature(color) -> tuple:
    try:
        color_type = int(color.type) if color.type is not None else None
    except Exception:
        color_type = str(getattr(color, "type", None))
    try:
        rgb = str(color.rgb) if color.rgb is not None else None
    except Exception:
        rgb = None
    try:
        theme = int(color.theme_color) if color.theme_color is not None else None
    except Exception:
        theme = None
    try:
        brightness = round(color.brightness, 5)
    except Exception:
        brightness = None
    return (color_type, rgb, theme, brightness)


def font_signature(font) -> tuple:
    try:
        language_value = font.language_id
        language = int(language_value) if language_value is not None else None
    except Exception:
        # 上游 Google Slides 文稿存在 `lang="en"`，它不在
        # python-pptx 的 MSO_LANGUAGE_ID 枚举映射中，直接保留 XML 原值。
        language = font._rPr.get("lang") if font._rPr is not None else None
    try:
        hyperlink = font._parent.hyperlink.address
    except Exception:
        hyperlink = None
    return (
        font.name,
        round(font.size.pt, 4) if font.size else None,
        bool(font.bold),
        bool(font.italic),
        bool(font.underline),
        language,
        color_signature(font.color),
        hyperlink,
    )


def text_signature(shape) -> tuple:
    text_frame = shape.text_frame
    paragraphs = []
    for paragraph in text_frame.paragraphs:
        runs = tuple(
            (run.text, font_signature(run.font))
            for run in paragraph.runs
        )
        paragraphs.append(
            (
                paragraph.text,
                paragraph.level,
                str(paragraph.alignment),
                str(paragraph.line_spacing),
                str(paragraph.space_before),
                str(paragraph.space_after),
                runs,
            )
        )
    return (
        (
            text_frame.margin_left,
            text_frame.margin_right,
            text_frame.margin_top,
            text_frame.margin_bottom,
            str(text_frame.vertical_anchor),
            text_frame.word_wrap,
            str(text_frame.auto_size),
        ),
        tuple(paragraphs),
    )


def shape_signature(shape) -> tuple:
    text = text_signature(shape) if getattr(shape, "has_text_frame", False) else ()
    picture = ()
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        try:
            picture_hash = hashlib.sha256(shape.image.blob).hexdigest()
        except Exception as exc:
            picture_hash = f"unreadable:{type(exc).__name__}"
        picture = (
            picture_hash,
            round(float(shape.crop_left), 6),
            round(float(shape.crop_right), 6),
            round(float(shape.crop_top), 6),
            round(float(shape.crop_bottom), 6),
        )
    children = ()
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        children = tuple(shape_signature(child) for child in shape.shapes)
    return (
        int(shape.shape_id),
        shape.name,
        str(shape.shape_type),
        placeholder_type(shape),
        rounded_geometry(shape),
        round(float(shape.rotation or 0), 4),
        text,
        picture,
        children,
    )


def slide_signature(slide) -> tuple:
    try:
        layout = str(slide.slide_layout.part.partname)
    except Exception:
        layout = ""
    return (
        slide._element.get("show"),
        layout,
        tuple(shape_signature(shape) for shape in slide.shapes),
    )


def notes_signature(slide) -> tuple:
    if not slide.has_notes_slide:
        return ()
    return tuple(shape_signature(shape) for shape in slide.notes_slide.shapes)


def local_name(element_or_name) -> str:
    value = element_or_name.tag if hasattr(element_or_name, "tag") else element_or_name
    return etree.QName(value).localname


def slide_roots(path: Path) -> list:
    with zipfile.ZipFile(path) as archive:
        return [
            etree.fromstring(archive.read(f"ppt/slides/slide{index}.xml"))
            for index in range(1, 10)
        ]


def transition_inventory(path: Path) -> tuple[list[dict], bool, bool]:
    inventory: list[dict] = []
    all_fade = True
    all_safe = True
    for slide_number, root in enumerate(slide_roots(path), start=1):
        transitions = root.xpath("./p:transition", namespaces={"p": P_NS})
        item = {
            "slide": slide_number,
            "count": len(transitions),
            "effect_children": [],
            "attributes": {},
            "fade_attributes": {},
            "fade": False,
            "safe": False,
            "issues": [],
        }
        if len(transitions) != 1:
            item["issues"].append("转场节点数不是 1")
            all_fade = False
            all_safe = False
            inventory.append(item)
            continue
        transition = transitions[0]
        attributes = {local_name(key): value for key, value in transition.attrib.items()}
        item["attributes"] = attributes
        effect_children = [
            local_name(child)
            for child in transition
            if local_name(child) != "extLst"
        ]
        item["effect_children"] = effect_children
        fade_nodes = [child for child in transition if local_name(child) == "fade"]
        item["fade"] = effect_children == ["fade"] and len(fade_nodes) == 1
        if not item["fade"]:
            item["issues"].append("直接转场效果不是唯一 fade")
            all_fade = False

        allowed_transition_attributes = {"spd", "advClick", "dur"}
        unknown_transition_attributes = sorted(
            set(attributes) - allowed_transition_attributes
        )
        if unknown_transition_attributes:
            item["issues"].append(
                f"未允许的 transition 属性={unknown_transition_attributes}"
            )
        if "advTm" in attributes:
            item["issues"].append("存在自动换页 advTm")
        if attributes.get("advClick", "1").lower() not in {"1", "true"}:
            item["issues"].append("禁用了鼠标单击换页")
        if attributes.get("spd", "med") not in {"slow", "med", "fast"}:
            item["issues"].append(f"不合理速度={attributes.get('spd')}")
        if "dur" in attributes:
            try:
                duration = int(attributes["dur"])
                if not 50 <= duration <= 10_000:
                    item["issues"].append(f"不合理时长={duration}")
            except Exception:
                item["issues"].append(f"时长非整数={attributes['dur']}")

        if fade_nodes:
            fade_attributes = {
                local_name(key): value for key, value in fade_nodes[0].attrib.items()
            }
            item["fade_attributes"] = fade_attributes
            unknown_fade_attributes = sorted(set(fade_attributes) - {"thruBlk"})
            if unknown_fade_attributes:
                item["issues"].append(
                    f"未允许的 fade 属性={unknown_fade_attributes}"
                )
            if fade_attributes.get("thruBlk", "0").lower() not in {
                "0",
                "1",
                "false",
                "true",
            }:
                item["issues"].append(
                    f"非法 Through Black 取值={fade_attributes.get('thruBlk')}"
                )

        sound_nodes = [
            local_name(node)
            for node in transition.iterdescendants()
            if local_name(node) in {"sndAc", "stSnd", "endSnd", "snd"}
        ]
        if sound_nodes:
            item["issues"].append(f"转场声音={sound_nodes}")
        item["safe"] = item["fade"] and not item["issues"]
        if not item["safe"]:
            all_safe = False
        inventory.append(item)
    return inventory, all_fade, all_safe


def timing_signatures(path: Path) -> list[tuple[bytes, ...]]:
    result = []
    for root in slide_roots(path):
        nodes = root.xpath("./p:timing", namespaces={"p": P_NS})
        result.append(tuple(etree.tostring(node, method="c14n") for node in nodes))
    return result


def content_slide_hashes(path: Path) -> list[str]:
    hashes = []
    for root in slide_roots(path):
        nodes = root.xpath("./p:cSld", namespaces={"p": P_NS})
        canonical = etree.tostring(nodes[0], method="c14n") if len(nodes) == 1 else b""
        hashes.append(hashlib.sha256(canonical).hexdigest())
    return hashes


def media_hashes(path: Path) -> Counter[str]:
    with zipfile.ZipFile(path) as archive:
        return Counter(
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("ppt/media/")
        )


def stable_member_names(path: Path) -> set[str]:
    with zipfile.ZipFile(path) as archive:
        return {
            name
            for name in archive.namelist()
            if not name.endswith("/")
            and not name.startswith("docProps/")
            and not re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
        }


def critical_hashes(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        return {
            name: hashlib.sha256(archive.read(name)).hexdigest()
            for name in sorted(stable_member_names(path))
        }


def package_ok(path: Path) -> tuple[bool, str]:
    try:
        with zipfile.ZipFile(path) as archive:
            broken = archive.testzip()
            names = set(archive.namelist())
            required = {
                "[Content_Types].xml",
                "ppt/presentation.xml",
                "ppt/slides/slide1.xml",
                "ppt/slides/slide9.xml",
            }
            missing = sorted(required - names)
            macros = sorted(
                name for name in names if name.lower().endswith("vbaproject.bin")
            )
        return broken is None and not missing and not macros, (
            f"损坏成员={broken}，缺失={missing}，宏={macros}"
        )
    except Exception as exc:
        return False, f"ZIP 检查失败：{exc}"


def full_slide_pictures(presentation: Presentation) -> list[str]:
    values: list[str] = []
    page_area = presentation.slide_width * presentation.slide_height
    for slide_index, slide in enumerate(presentation.slides, start=1):
        for shape in iter_shapes(slide.shapes):
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            coverage = shape.width * shape.height / page_area
            if coverage >= 0.85:
                values.append(f"slide{slide_index}:{shape.name}:{coverage:.3f}")
    return values


def render_presentation(path: Path, outdir: Path) -> tuple[list[Path], int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "lo-profile"
    profile.mkdir()
    completed = subprocess.run(
        [
            "soffice",
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(outdir),
            str(path),
        ],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return [], 0, f"转 PDF 失败：{completed.stdout[-500:]}"
    info = subprocess.run(
        ["pdfinfo", str(pdf)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=30,
        check=False,
    ).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    pages = int(match.group(1)) if match else 0
    prefix = outdir / "slide"
    raster = subprocess.run(
        ["pdftoppm", "-jpeg", "-r", "90", str(pdf), str(prefix)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    images = sorted(outdir.glob("slide-*.jpg")) if raster.returncode == 0 else []
    return images, pages, (
        f"页数={pages}，渲染图={len(images)}，"
        f"soffice={completed.returncode}，pdftoppm={raster.returncode}"
    )


def image_similarity(first: Path, second: Path) -> float:
    with Image.open(first) as left_image, Image.open(second) as right_image:
        left = left_image.convert("RGB")
        right = right_image.convert("RGB")
        if left.size != right.size:
            return 0.0
        difference = ImageChops.difference(left, right)
        histogram = difference.histogram()
        squared = sum(
            count * ((index % 256) ** 2)
            for index, count in enumerate(histogram)
        )
        pixels = left.width * left.height
        rms = math.sqrt(squared / (pixels * 3)) if pixels else 255.0
        return max(0.0, 1.0 - rms / 255.0)


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {
        "criterion_id": identifier,
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
        "raw": {
            "description": description,
            "weight": weight,
            "type": kind,
            "passed": bool(passed),
        },
    }


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [
        item["criterion_id"]
        for item in criteria
        if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]
    ]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(
        item["raw"]["weight"]
        for item in criteria
        if item["raw"]["weight"] > 0 and item["raw"]["passed"]
    )
    penalty = sum(
        abs(item["raw"]["weight"])
        for item in criteria
        if item["raw"]["weight"] < 0 and not item["raw"]["passed"]
    )
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [sha256(path)] if path.is_file() else [],
        "criterion_results": criteria,
        "reward": reward,
        "passed": reward >= 0.8 and not hurdles,
        "hurdle_failures": hurdles,
    }


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        criteria = [
            result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配")
            for identifier in SPECS
        ]
        return make_result(path, criteria)
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        criteria = [
            result_item(identifier, False, f"PPTX 打开失败：{exc}")
            for identifier in SPECS
        ]
        return make_result(path, criteria)

    criteria: list[dict] = []
    package_passed, package_evidence = package_ok(path)
    size_passed = (
        len(source.slides) == 9
        and len(candidate.slides) == 9
        and candidate.slide_width == source.slide_width
        and candidate.slide_height == source.slide_height
    )
    criteria.append(
        result_item(
            "R001",
            package_passed and size_passed,
            f"{package_evidence}，输入页数={len(source.slides)}，"
            f"候选页数={len(candidate.slides)}，页面="
            f"{candidate.slide_width / EMU_PER_INCH:.4f}×"
            f"{candidate.slide_height / EMU_PER_INCH:.4f}英寸",
        )
    )

    try:
        transitions, all_fade, all_safe = transition_inventory(path)
    except Exception as exc:
        transitions, all_fade, all_safe = [], False, False
        transition_error = str(exc)
    else:
        transition_error = ""
    summary = [
        {
            "slide": item["slide"],
            "count": item["count"],
            "effects": item["effect_children"],
            "attributes": item["attributes"],
            "fade_attributes": item["fade_attributes"],
            "issues": item["issues"],
        }
        for item in transitions
    ]
    criteria.append(
        result_item(
            "R002",
            len(transitions) == 9 and all_fade,
            f"逐页转场={json.dumps(summary, ensure_ascii=False)}，解析错误={transition_error}",
        )
    )
    criteria.append(
        result_item(
            "R003",
            len(transitions) == 9 and all_safe,
            f"速度/时长/换页/声音检查通过={all_safe}，"
            f"逐页问题={[item['issues'] for item in transitions]}",
        )
    )

    semantic_mismatches: list[int] = []
    notes_mismatches: list[int] = []
    if len(candidate.slides) == 9:
        for index in range(9):
            if slide_signature(candidate.slides[index]) != slide_signature(source.slides[index]):
                semantic_mismatches.append(index + 1)
            if notes_signature(candidate.slides[index]) != notes_signature(source.slides[index]):
                notes_mismatches.append(index + 1)
    else:
        semantic_mismatches = list(range(1, 10))
        notes_mismatches = list(range(1, 10))
    criteria.append(
        result_item(
            "R004",
            not semantic_mismatches,
            f"内容/样式/对象/图片/几何/顺序/版式变化页={semantic_mismatches}",
        )
    )

    try:
        timing_mismatches = [
            index + 1
            for index, (expected, actual) in enumerate(
                zip(
                    timing_signatures(source_path),
                    timing_signatures(path),
                    strict=True,
                )
            )
            if expected != actual
        ]
    except Exception:
        timing_mismatches = list(range(1, 10))
    criteria.append(
        result_item(
            "R005",
            not notes_mismatches and not timing_mismatches,
            f"备注变化页={notes_mismatches}，动画/时序变化页={timing_mismatches}",
        )
    )

    try:
        source_media = media_hashes(source_path)
        candidate_media = media_hashes(path)
        media_equal = source_media == candidate_media
        source_members = stable_member_names(source_path)
        candidate_members = stable_member_names(path)
        member_sets_equal = source_members == candidate_members
    except Exception:
        source_media = Counter()
        candidate_media = Counter()
        media_equal = False
        member_sets_equal = False
        source_members = set()
        candidate_members = set()
    flat_pictures = full_slide_pictures(candidate)
    structure_passed = media_equal and member_sets_equal and not flat_pictures
    criteria.append(
        result_item(
            "R006",
            structure_passed,
            f"媒体哈希相等={media_equal}，媒体数="
            f"{sum(source_media.values())}→{sum(candidate_media.values())}，"
            f"包成员集相等={member_sets_equal}，整页图片={flat_pictures}，"
            f"缺失成员={sorted(source_members - candidate_members)}，"
            f"额外成员={sorted(candidate_members - source_members)}",
        )
    )

    native_passed = (
        package_passed
        and len(transitions) == 9
        and all_fade
        and not flat_pictures
    )
    criteria.append(
        result_item(
            "R008",
            native_passed,
            f"包完整={package_passed}，幻灯片级 Fade={all_fade}，"
            f"转场记录数={len(transitions)}，整页图片={flat_pictures}",
        )
    )

    try:
        source_content_hashes = content_slide_hashes(source_path)
        candidate_content_hashes = content_slide_hashes(path)
        exact_content_pages = [
            index + 1
            for index, (expected, actual) in enumerate(
                zip(source_content_hashes, candidate_content_hashes, strict=True)
            )
            if expected == actual
        ]
        critical_equal = critical_hashes(source_path) == critical_hashes(path)
    except Exception:
        exact_content_pages = []
        critical_equal = False
    exact_package_passed = exact_content_pages == list(range(1, 10)) and critical_equal
    criteria.append(
        result_item(
            "R009",
            exact_package_passed,
            f"去除转场后 cSld 精确保持页={exact_content_pages}，"
            f"关键非幻灯片包部件哈希保持={critical_equal}",
        )
    )

    artifact_files = sorted(
        str(item.relative_to(artifact_dir))
        for item in artifact_dir.rglob("*")
        if item.is_file()
    )
    all_text = "\n".join(
        shape.text
        for slide in candidate.slides
        for shape in iter_shapes(slide.shapes)
        if getattr(shape, "has_text_frame", False)
    )
    all_notes = "\n".join(
        shape.text
        for slide in candidate.slides
        if slide.has_notes_slide
        for shape in slide.notes_slide.shapes
        if getattr(shape, "has_text_frame", False)
    )
    lowered = f"{all_text}\n{all_notes}".lower()
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker.lower() in lowered]
    hygiene_passed = artifact_files == [OUTPUT_NAME] and not markers
    criteria.append(
        result_item(
            "R010",
            hygiene_passed,
            f"输出文件={artifact_files}，可疑标记={markers}",
        )
    )

    pre_render_ready = (
        package_passed
        and size_passed
        and len(transitions) == 9
        and all_fade
        and all_safe
        and not semantic_mismatches
        and not notes_mismatches
        and not timing_mismatches
        and structure_passed
        and native_passed
    )
    visual_passed = False
    visual_min = 0.0
    render_evidence = "前置结构、转场或保持检查失败，未执行渲染"
    if pre_render_ready:
        with tempfile.TemporaryDirectory(prefix="pptarena037-judge-") as temporary:
            root = Path(temporary)
            source_images, source_pages, source_note = render_presentation(
                source_path, root / "source"
            )
            candidate_images, candidate_pages, candidate_note = render_presentation(
                path, root / "candidate"
            )
            render_evidence = f"源[{source_note}]；候选[{candidate_note}]"
            if len(source_images) == 9 and len(candidate_images) == 9:
                scores = [
                    image_similarity(expected, actual)
                    for expected, actual in zip(source_images, candidate_images, strict=True)
                ]
                visual_min = min(scores)
            visual_passed = (
                source_pages == 9
                and candidate_pages == 9
                and len(source_images) == 9
                and len(candidate_images) == 9
                and visual_min >= 0.9985
            )
    criteria.append(
        result_item(
            "R007",
            visual_passed,
            f"9 页最低像素相似度={visual_min:.6f}；{render_evidence}",
        )
    )

    severe = (
        not pre_render_ready
        or not visual_passed
        or bool(flat_pictures)
        or bool(markers)
    )
    criteria.append(result_item("R011", not severe, f"严重错误={severe}"))
    order = {identifier: index for index, identifier in enumerate(SPECS)}
    criteria.sort(key=lambda item: order[item["criterion_id"]])
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    output = args.artifact_dir / OUTPUT_NAME
    if output.is_file():
        evaluation = evaluate(output, args.artifact_dir)
    else:
        evaluation = {
            "task_id": TASK_ID,
            "rubric_version": RUBRIC_VERSION,
            "judge_version": JUDGE_VERSION,
            "artifact_sha256": [],
            "criterion_results": [],
            "reward": 0.0,
            "passed": False,
            "hurdle_failures": ["missing_output"],
        }
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(
        json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
