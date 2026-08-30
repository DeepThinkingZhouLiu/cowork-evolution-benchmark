#!/usr/bin/env python3
"""只为 PPTArena Case 41 构造可重放的合法变体、边界例和负例。"""

from __future__ import annotations

import argparse
import shutil
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.util import Inches


OUTPUT_NAME = "MatchBackgroundtoPicture_GroundTruthA.pptx"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def qn(namespace: str, local: str) -> str:
    return f"{{{namespace}}}{local}"


def output_path(root: Path, case: str) -> Path:
    target = root / case / OUTPUT_NAME
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def serialize_xml(root) -> bytes:
    return etree.tostring(
        root,
        encoding="UTF-8",
        xml_declaration=True,
        standalone=True,
    )


def rewrite_archive(source: Path, output: Path, transform) -> None:
    with zipfile.ZipFile(source) as archive_in, zipfile.ZipFile(output, "w") as archive_out:
        for info in archive_in.infolist():
            data = archive_in.read(info.filename)
            archive_out.writestr(info, transform(info.filename, data))


def replace_background(
    root,
    stops: list[tuple[int, str, str, list[tuple[str, int]]]],
    *,
    angle: int = 5_400_000,
    solid_rgb: str | None = None,
) -> None:
    c_sld = root.find(qn(P_NS, "cSld"))
    if c_sld is None:
        raise ValueError("缺失 p:cSld")
    for background in c_sld.findall(qn(P_NS, "bg")):
        c_sld.remove(background)
    background = etree.Element(qn(P_NS, "bg"))
    background_properties = etree.SubElement(background, qn(P_NS, "bgPr"))
    if solid_rgb is not None:
        solid = etree.SubElement(background_properties, qn(A_NS, "solidFill"))
        etree.SubElement(solid, qn(A_NS, "srgbClr"), val=solid_rgb)
    else:
        gradient = etree.SubElement(background_properties, qn(A_NS, "gradFill"), rotWithShape="1")
        stop_list = etree.SubElement(gradient, qn(A_NS, "gsLst"))
        for position, color_kind, value, transforms in stops:
            stop = etree.SubElement(stop_list, qn(A_NS, "gs"), pos=str(position))
            color = etree.SubElement(stop, qn(A_NS, color_kind), val=value)
            for transform, amount in transforms:
                etree.SubElement(color, qn(A_NS, transform), val=str(amount))
        etree.SubElement(gradient, qn(A_NS, "lin"), ang=str(angle), scaled="1")
    etree.SubElement(background_properties, qn(A_NS, "effectLst"))
    c_sld.insert(0, background)


def gradient_from_source(
    source: Path,
    output: Path,
    stops: list[tuple[int, str, str, list[tuple[str, int]]]],
    *,
    angle: int = 5_400_000,
) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name != "ppt/slides/slide1.xml":
            return data
        root = etree.fromstring(data)
        replace_background(root, stops, angle=angle)
        return serialize_xml(root)

    rewrite_archive(source, output, transform)


def solid_background(source: Path, output: Path, rgb: str) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name != "ppt/slides/slide1.xml":
            return data
        root = etree.fromstring(data)
        replace_background(root, [], solid_rgb=rgb)
        return serialize_xml(root)

    rewrite_archive(source, output, transform)


def add_picture_border(
    source: Path,
    output: Path,
    *,
    width_emu: int,
    kind: str,
    value: str,
) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name != "ppt/slides/slide1.xml":
            return data
        root = etree.fromstring(data)
        pictures = root.xpath("./p:cSld/p:spTree/p:pic", namespaces={"p": P_NS})
        if len(pictures) != 1:
            raise ValueError("中央图片不唯一")
        properties = pictures[0].find(qn(P_NS, "spPr"))
        if properties is None:
            raise ValueError("图片缺失 p:spPr")
        for line in properties.findall(qn(A_NS, "ln")):
            properties.remove(line)
        line = etree.SubElement(properties, qn(A_NS, "ln"), w=str(width_emu))
        fill = etree.SubElement(line, qn(A_NS, "solidFill"))
        etree.SubElement(fill, qn(A_NS, kind), val=value)
        etree.SubElement(line, qn(A_NS, "prstDash"), val="solid")
        return serialize_xml(root)

    rewrite_archive(source, output, transform)


def repacked_content_types(golden: Path, output: Path) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name == "[Content_Types].xml":
            return data.replace(b"</Types>", b"  \n</Types>")
        return data

    rewrite_archive(golden, output, transform)


def changed_title(golden: Path, output: Path) -> None:
    changed = False

    def transform(name: str, data: bytes) -> bytes:
        nonlocal changed
        if name != "ppt/slides/slide1.xml":
            return data
        root = etree.fromstring(data)
        nodes = root.xpath(".//a:t", namespaces={"a": A_NS})
        if not nodes:
            raise ValueError("无法找到标题文字")
        nodes[0].text = f"{nodes[0].text} X"
        changed = True
        return serialize_xml(root)

    rewrite_archive(golden, output, transform)
    if not changed:
        raise ValueError("未改写标题")


def moved_picture(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    picture = next(
        shape
        for shape in presentation.slides[0].shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    )
    picture.left += Inches(0.25)
    presentation.save(output)


def changed_crop(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    picture = next(
        shape
        for shape in presentation.slides[0].shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    )
    picture.crop_left = 0.03
    presentation.save(output)


def missing_picture(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    picture = next(
        shape
        for shape in presentation.slides[0].shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    )
    picture._element.getparent().remove(picture._element)
    presentation.save(output)


def background_rectangle(source: Path, output: Path) -> None:
    presentation = Presentation(source)
    slide = presentation.slides[0]
    rectangle = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        0,
        0,
        presentation.slide_width,
        presentation.slide_height,
    )
    rectangle.fill.solid()
    rectangle.fill.fore_color.rgb = RGBColor(224, 232, 240)
    rectangle.line.fill.background()
    # 即使移到最底层，它仍是图形伪背景，不是原生幻灯片背景。
    tree = slide.shapes._spTree
    tree.remove(rectangle._element)
    tree.insert(2, rectangle._element)
    presentation.save(output)


def hardcoded_marker(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    box = presentation.slides[0].shapes.add_textbox(
        Inches(0.1), Inches(0.1), Inches(2.5), Inches(0.25)
    )
    box.text = "solution/Golden"
    presentation.save(output)


def extra_slide(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    presentation.slides.add_slide(presentation.slide_layouts[0])
    presentation.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--upstream-unfit", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()

    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))
    shutil.copy2(args.upstream_unfit, output_path(args.outdir, "upstream-unfit"))

    # 独立合法变体：两个端点色直接来自原图浅色像素，不复制 Golden 的主题色参数。
    gradient_from_source(
        args.source,
        output_path(args.outdir, "valid-rgb-sampled"),
        [
            (0, "srgbClr", "FEFEFE", []),
            (100_000, "srgbClr", "B0C0D8", []),
        ],
    )
    add_picture_border(
        args.golden,
        output_path(args.outdir, "valid-thin-accent-border"),
        width_emu=12_700,
        kind="schemeClr",
        value="accent1",
    )
    repacked_content_types(
        args.golden,
        output_path(args.outdir, "repacked-content-types"),
    )

    solid_background(args.source, output_path(args.outdir, "solid-gray"), "E0E8F0")
    gradient_from_source(
        args.source,
        output_path(args.outdir, "horizontal-gradient"),
        [(0, "srgbClr", "FEFEFE", []), (100_000, "srgbClr", "B0C0D8", [])],
        angle=0,
    )
    gradient_from_source(
        args.source,
        output_path(args.outdir, "pure-white-gradient"),
        [(0, "srgbClr", "FFFFFF", []), (100_000, "srgbClr", "FAFAFA", [])],
    )
    gradient_from_source(
        args.source,
        output_path(args.outdir, "too-dark-bottom"),
        [(0, "srgbClr", "FEFEFE", []), (100_000, "srgbClr", "52677F", [])],
    )
    gradient_from_source(
        args.source,
        output_path(args.outdir, "yellow-gradient"),
        [(0, "srgbClr", "FFFFFF", []), (100_000, "srgbClr", "F0D77A", [])],
    )
    gradient_from_source(
        args.source,
        output_path(args.outdir, "reversed-gradient"),
        [(0, "srgbClr", "B0C0D8", []), (100_000, "srgbClr", "FEFEFE", [])],
    )
    background_rectangle(args.source, output_path(args.outdir, "background-rectangle"))
    changed_title(args.golden, output_path(args.outdir, "changed-title"))
    moved_picture(args.golden, output_path(args.outdir, "moved-picture"))
    changed_crop(args.golden, output_path(args.outdir, "changed-crop"))
    add_picture_border(
        args.golden,
        output_path(args.outdir, "thick-red-border"),
        width_emu=76_200,
        kind="srgbClr",
        value="FF0000",
    )
    missing_picture(args.golden, output_path(args.outdir, "missing-picture"))
    hardcoded_marker(args.golden, output_path(args.outdir, "hardcoded-marker"))
    extra_slide(args.golden, output_path(args.outdir, "extra-slide"))


if __name__ == "__main__":
    main()
