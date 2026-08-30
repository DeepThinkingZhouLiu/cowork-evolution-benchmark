#!/usr/bin/env python3
"""PPTArena Case 41 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import re
import statistics
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from lxml import etree
from PIL import Image, ImageChops
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "MatchBackgroundtoPicture_GroundTruthA.pptx"
INPUT_NAME = "MatchBackgroundtoPicture_TestA.pptx"
INPUT_SHA256 = "8c789d83d7063549ab3a198a7363e7e6d9feda84ea20865ae00a6307d3be52cd"
TASK_ID = "cowork-evo/pptarena-041-picture-matched-gradient-background"
RUBRIC_VERSION = "pptarena-case-041-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-041-deterministic-v0.1"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
EMU_PER_INCH = 914_400
SUSPICIOUS_MARKERS = [
    "solution/golden",
    "/root/",
    "judge-result",
    "artifact_dir",
    "pptarena-case-041",
]

SPECS = {
    "R001": ("指定 PPTX 可打开，页面尺寸不变且总页数为 1", 6, "hurdle"),
    "R002": ("第 1 页使用至少两端色标的原生纵向线性渐变背景", 18, "hurdle"),
    "R003": ("渐变从近白到明亮浅蓝灰，来自 accent1 亮度变体或原图浅色像素", 16, "hurdle"),
    "R004": ("标题、页脚、图片、样式、几何、链接和顺序保持", 18, "hurdle"),
    "R005": ("中央 TIFF 字节、裁剪和几何不变，可选边框细且为蓝色", 10, "hurdle"),
    "R006": ("页面可渲染，中央热力图渲染内容与冻结输入保持", 10, "hurdle"),
    "R007": ("实际背景从上到下逐渐加深，两端明亮并与原图协调", 8, "hurdle"),
    "R008": ("保持 3 个原生页面对象、1 个原媒体和原版式关系，无压平、宏或额外素材", 7, "hurdle"),
    "R009": ("输出目录干净且不含答案路径或未授权文字", 5, "positive"),
    "R010": ("除背景和合法可选边框外，内容 XML 和关键包部件字节级保持", 2, "positive"),
    "R011": ("缺少合法渐变、中央图片错渲染、原内容变化或可编辑性破坏", -20, "penalty"),
}


def normalize_visible(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def locate_source() -> Path | None:
    candidates = [
        Path("/app/input_files") / INPUT_NAME,
        Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME,
    ]
    return next((path for path in candidates if path.is_file()), None)


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def placeholder_type(shape) -> str:
    if not getattr(shape, "is_placeholder", False):
        return ""
    return str(shape.placeholder_format.type)


def rounded_geometry(shape) -> tuple:
    return tuple(
        round(int(value) / 1000) * 1000
        for value in (shape.left, shape.top, shape.width, shape.height)
    )


def color_signature(color) -> tuple:
    try:
        color_type = int(color.type) if color.type is not None else None
    except Exception:
        color_type = str(getattr(color, "type", None))
    try:
        rgb = str(color.rgb) if color.rgb is not None else None
    except Exception:
        rgb = None
    try:
        theme = int(color.theme_color) if color.theme_color is not None else None
    except Exception:
        theme = None
    try:
        brightness = round(color.brightness, 5)
    except Exception:
        brightness = None
    return (color_type, rgb, theme, brightness)


def font_signature(font, run=None) -> tuple:
    try:
        language_value = font.language_id
        language = int(language_value) if language_value is not None else None
    except Exception:
        language = font._rPr.get("lang") if font._rPr is not None else None
    try:
        hyperlink = run.hyperlink.address if run is not None else None
    except Exception:
        hyperlink = None
    return (
        font.name,
        round(font.size.pt, 4) if font.size else None,
        bool(font.bold),
        bool(font.italic),
        bool(font.underline),
        language,
        color_signature(font.color),
        hyperlink,
    )


def text_signature(shape) -> tuple:
    text_frame = shape.text_frame
    paragraphs = []
    for paragraph in text_frame.paragraphs:
        runs = tuple(
            (run.text, font_signature(run.font, run))
            for run in paragraph.runs
        )
        paragraphs.append(
            (
                paragraph.text,
                paragraph.level,
                str(paragraph.alignment),
                str(paragraph.line_spacing),
                str(paragraph.space_before),
                str(paragraph.space_after),
                runs,
            )
        )
    return (
        (
            text_frame.margin_left,
            text_frame.margin_right,
            text_frame.margin_top,
            text_frame.margin_bottom,
            str(text_frame.vertical_anchor),
            text_frame.word_wrap,
            str(text_frame.auto_size),
        ),
        tuple(paragraphs),
    )


def picture_blob_hash(shape) -> str:
    try:
        return hashlib.sha256(shape.image.blob).hexdigest()
    except Exception as exc:
        return f"unreadable:{type(exc).__name__}"


def shape_signature(shape) -> tuple:
    text = text_signature(shape) if getattr(shape, "has_text_frame", False) else ()
    picture = ()
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        picture = (
            picture_blob_hash(shape),
            round(float(shape.crop_left), 6),
            round(float(shape.crop_right), 6),
            round(float(shape.crop_top), 6),
            round(float(shape.crop_bottom), 6),
        )
    children = ()
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        children = tuple(shape_signature(child) for child in shape.shapes)
    return (
        int(shape.shape_id),
        shape.name,
        str(shape.shape_type),
        placeholder_type(shape),
        rounded_geometry(shape),
        round(float(shape.rotation or 0), 4),
        text,
        picture,
        children,
    )


def slide_signature(slide) -> tuple:
    try:
        layout = str(slide.slide_layout.part.partname)
    except Exception:
        layout = ""
    return (
        slide._element.get("show"),
        layout,
        tuple(shape_signature(shape) for shape in slide.shapes),
    )


def qn(namespace: str, local: str) -> str:
    return f"{{{namespace}}}{local}"


def local_name(element_or_name) -> str:
    value = element_or_name.tag if hasattr(element_or_name, "tag") else element_or_name
    return etree.QName(value).localname


def slide_root(path: Path):
    with zipfile.ZipFile(path) as archive:
        return etree.fromstring(archive.read("ppt/slides/slide1.xml"))


def theme_accent1(path: Path) -> tuple[int, int, int]:
    with zipfile.ZipFile(path) as archive:
        theme_names = sorted(
            name for name in archive.namelist() if name.startswith("ppt/theme/theme") and name.endswith(".xml")
        )
        if not theme_names:
            raise ValueError("缺失主题颜色表")
        root = etree.fromstring(archive.read(theme_names[0]))
    nodes = root.xpath(
        ".//a:themeElements/a:clrScheme/a:accent1/*[1]",
        namespaces={"a": A_NS},
    )
    if len(nodes) != 1:
        raise ValueError("主题 accent1 不唯一")
    node = nodes[0]
    value = node.get("val") if local_name(node) == "srgbClr" else node.get("lastClr")
    if not value or not re.fullmatch(r"[0-9A-Fa-f]{6}", value):
        raise ValueError(f"无法解析 accent1={value}")
    return tuple(int(value[index : index + 2], 16) for index in (0, 2, 4))


def apply_color_transforms(
    rgb: tuple[int, int, int],
    color_node,
) -> tuple[int, int, int]:
    values = [float(value) for value in rgb]
    for transform in color_node:
        name = local_name(transform)
        try:
            amount = int(transform.get("val")) / 100_000
        except Exception:
            continue
        if name in {"lumMod", "shade"}:
            values = [value * amount for value in values]
        elif name == "lumOff":
            values = [value + 255 * amount for value in values]
        elif name == "tint":
            values = [value + (255 - value) * amount for value in values]
    return tuple(max(0, min(255, round(value))) for value in values)


def resolve_stop_color(color_node, accent_rgb: tuple[int, int, int]) -> dict:
    kind = local_name(color_node)
    value = color_node.get("val", "")
    base = None
    if kind == "srgbClr" and re.fullmatch(r"[0-9A-Fa-f]{6}", value):
        base = tuple(int(value[index : index + 2], 16) for index in (0, 2, 4))
    elif kind == "schemeClr" and value == "accent1":
        base = accent_rgb
    rgb = apply_color_transforms(base, color_node) if base is not None else None
    return {
        "kind": kind,
        "value": value,
        "rgb": list(rgb) if rgb is not None else None,
        "transforms": [
            {"name": local_name(child), "val": child.get("val")}
            for child in color_node
        ],
    }


def background_inventory(path: Path, accent_rgb: tuple[int, int, int]) -> dict:
    root = slide_root(path)
    backgrounds = root.xpath("./p:cSld/p:bg", namespaces={"p": P_NS})
    gradients = root.xpath(
        "./p:cSld/p:bg/p:bgPr/a:gradFill",
        namespaces={"p": P_NS, "a": A_NS},
    )
    evidence = {
        "background_count": len(backgrounds),
        "gradient_count": len(gradients),
        "angle_degrees": None,
        "vertical": False,
        "stops": [],
        "structure_passed": False,
        "color_passed": False,
        "top_rgb": None,
        "bottom_rgb": None,
        "theme_mode": False,
        "sampled_rgb_mode": False,
    }
    if len(backgrounds) != 1 or len(gradients) != 1:
        return evidence
    gradient = gradients[0]
    if gradient.xpath("./a:path", namespaces={"a": A_NS}):
        return evidence
    linear = gradient.xpath("./a:lin", namespaces={"a": A_NS})
    if len(linear) != 1:
        return evidence
    try:
        angle = int(linear[0].get("ang")) / 60_000 % 360
    except Exception:
        return evidence
    vertical = 75 <= angle <= 105 or 255 <= angle <= 285
    evidence["angle_degrees"] = angle
    evidence["vertical"] = vertical
    stops = []
    for stop in gradient.xpath("./a:gsLst/a:gs", namespaces={"a": A_NS}):
        try:
            position = int(stop.get("pos"))
        except Exception:
            continue
        color_nodes = [child for child in stop if local_name(child) in {"srgbClr", "schemeClr"}]
        if len(color_nodes) != 1:
            continue
        item = {"position": position, **resolve_stop_color(color_nodes[0], accent_rgb)}
        stops.append(item)
    stops.sort(key=lambda item: item["position"])
    evidence["stops"] = stops
    endpoints = (
        2 <= len(stops) <= 8
        and stops[0]["position"] <= 5_000
        and stops[-1]["position"] >= 95_000
        and all(item["rgb"] is not None for item in stops)
    )
    evidence["structure_passed"] = vertical and endpoints
    if not evidence["structure_passed"]:
        return evidence
    if 75 <= angle <= 105:
        top, bottom = tuple(stops[0]["rgb"]), tuple(stops[-1]["rgb"])
    else:
        top, bottom = tuple(stops[-1]["rgb"]), tuple(stops[0]["rgb"])
    evidence["top_rgb"] = list(top)
    evidence["bottom_rgb"] = list(bottom)
    evidence["theme_mode"] = all(
        item["kind"] == "schemeClr" and item["value"] == "accent1"
        for item in stops
    )
    evidence["sampled_rgb_mode"] = all(item["kind"] == "srgbClr" for item in stops)
    top_mean = statistics.mean(top)
    bottom_mean = statistics.mean(bottom)
    evidence["color_passed"] = (
        (evidence["theme_mode"] or evidence["sampled_rgb_mode"])
        and top_mean >= 235
        and max(top) - min(top) <= 30
        and 180 <= bottom_mean <= 245
        and bottom[2] >= bottom[0] + 5
        and max(bottom) - min(bottom) <= 65
        and top_mean - bottom_mean >= 10
        and math.dist(top, bottom) >= 15
    )
    return evidence


def source_picture_palette(source: Presentation) -> set[tuple[int, int, int]]:
    picture = next(
        shape
        for shape in source.slides[0].shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    )
    image = Image.open(io.BytesIO(picture.image.blob)).convert("RGB")
    return {
        tuple(min(255, round(channel / 8) * 8) for channel in pixel)
        for pixel in image.getdata()
    }


def palette_distance(rgb: tuple[int, int, int], palette: set[tuple[int, int, int]]) -> float:
    return min(math.dist(rgb, item) for item in palette)


def border_evidence(picture_shape, accent_rgb: tuple[int, int, int]) -> tuple[bool, dict]:
    lines = picture_shape._element.xpath("./p:spPr/a:ln")
    if not lines:
        return True, {"present": False}
    if len(lines) != 1:
        return False, {"present": True, "count": len(lines)}
    line = lines[0]
    try:
        width = int(line.get("w", "12700"))
    except Exception:
        width = 10**9
    solid = line.xpath("./a:solidFill/*[1]")
    color = resolve_stop_color(solid[0], accent_rgb) if len(solid) == 1 else {
        "kind": "",
        "value": "",
        "rgb": None,
        "transforms": [],
    }
    rgb = tuple(color["rgb"]) if color["rgb"] is not None else None
    color_ok = (
        color["kind"] == "schemeClr" and color["value"] == "accent1"
        or rgb is not None
        and rgb[2] >= rgb[0]
        and math.dist(rgb, accent_rgb) <= 100
    )
    dash = line.xpath("./a:prstDash/@val")
    dash_ok = not dash or dash == ["solid"]
    no_fill = bool(line.xpath("./a:noFill"))
    passed = 0 < width <= 25_400 and color_ok and dash_ok and not no_fill
    return passed, {
        "present": True,
        "width_emu": width,
        "width_pt": round(width / 12_700, 4),
        "color": color,
        "dash": dash,
        "no_fill": no_fill,
    }


def media_hashes(path: Path) -> Counter[str]:
    with zipfile.ZipFile(path) as archive:
        return Counter(
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("ppt/media/")
        )


def stable_member_names(path: Path) -> set[str]:
    with zipfile.ZipFile(path) as archive:
        return {
            name
            for name in archive.namelist()
            if not name.endswith("/")
            and not name.startswith("docProps/")
            and name != "ppt/slides/slide1.xml"
        }


def critical_hashes(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        return {
            name: hashlib.sha256(archive.read(name)).hexdigest()
            for name in sorted(stable_member_names(path))
        }


def normalized_content_hash(path: Path, remove_picture_line: bool) -> str:
    root = slide_root(path)
    for background in root.xpath("./p:cSld/p:bg", namespaces={"p": P_NS}):
        background.getparent().remove(background)
    if remove_picture_line:
        for line in root.xpath("./p:cSld/p:spTree/p:pic/p:spPr/a:ln", namespaces={"p": P_NS, "a": A_NS}):
            line.getparent().remove(line)
    nodes = root.xpath("./p:cSld", namespaces={"p": P_NS})
    canonical = etree.tostring(nodes[0], method="c14n") if len(nodes) == 1 else b""
    return hashlib.sha256(canonical).hexdigest()


def package_ok(path: Path) -> tuple[bool, str]:
    try:
        with zipfile.ZipFile(path) as archive:
            broken = archive.testzip()
            names = set(archive.namelist())
            required = {
                "[Content_Types].xml",
                "ppt/presentation.xml",
                "ppt/slides/slide1.xml",
            }
            missing = sorted(required - names)
            macros = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
        return broken is None and not missing and not macros, (
            f"损坏成员={broken}，缺失={missing}，宏={macros}"
        )
    except Exception as exc:
        return False, f"ZIP 检查失败：{exc}"


def full_slide_pictures(presentation: Presentation) -> list[str]:
    values = []
    page_area = presentation.slide_width * presentation.slide_height
    for slide_index, slide in enumerate(presentation.slides, start=1):
        for shape in iter_shapes(slide.shapes):
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            coverage = shape.width * shape.height / page_area
            if coverage >= 0.85:
                values.append(f"slide{slide_index}:{shape.name}:{coverage:.3f}")
    return values


def render_presentation(path: Path, outdir: Path) -> tuple[list[Path], int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "lo-profile"
    profile.mkdir()
    completed = subprocess.run(
        [
            "soffice",
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(outdir),
            str(path),
        ],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return [], 0, f"转 PDF 失败：{completed.stdout[-500:]}"
    info = subprocess.run(
        ["pdfinfo", str(pdf)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=30,
        check=False,
    ).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    pages = int(match.group(1)) if match else 0
    prefix = outdir / "slide"
    raster = subprocess.run(
        ["pdftoppm", "-jpeg", "-r", "120", str(pdf), str(prefix)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    images = sorted(outdir.glob("slide-*.jpg")) if raster.returncode == 0 else []
    return images, pages, (
        f"页数={pages}，渲染图={len(images)}，"
        f"soffice={completed.returncode}，pdftoppm={raster.returncode}"
    )


def image_similarity(first: Image.Image, second: Image.Image) -> float:
    left = first.convert("RGB")
    right = second.convert("RGB")
    if left.size != right.size:
        return 0.0
    difference = ImageChops.difference(left, right)
    histogram = difference.histogram()
    squared = sum(count * ((index % 256) ** 2) for index, count in enumerate(histogram))
    pixels = left.width * left.height
    rms = math.sqrt(squared / (pixels * 3)) if pixels else 255.0
    return max(0.0, 1.0 - rms / 255.0)


def picture_crop(image: Image.Image, shape, slide_size: tuple[int, int]) -> Image.Image:
    slide_width, slide_height = slide_size
    left = round(shape.left * image.width / slide_width)
    top = round(shape.top * image.height / slide_height)
    right = round((shape.left + shape.width) * image.width / slide_width)
    bottom = round((shape.top + shape.height) * image.height / slide_height)
    # 任务允许最多 2pt 的可选细边框。LibreOffice 会对图片边界
    # 做抗锯齿重采样，因此只比较足够深的图片内部；字节、裁剪和几何
    # 仍由 R004/R005 独立严格保护。
    padding = max(3, round(min(image.size) * 0.012))
    return image.crop((left + padding, top + padding, right - padding, bottom - padding))


def side_band_color(image: Image.Image, y_ratio: float) -> tuple[int, int, int]:
    rgb = image.convert("RGB")
    y = max(0, min(rgb.height - 1, round(rgb.height * y_ratio)))
    edge = max(8, round(rgb.width * 0.03))
    pixels = [rgb.getpixel((x, y)) for x in range(4, edge)]
    pixels += [rgb.getpixel((x, y)) for x in range(rgb.width - edge, rgb.width - 4)]
    return tuple(round(statistics.median(pixel[index] for pixel in pixels)) for index in range(3))


def rendered_background_evidence(
    candidate_image: Image.Image,
    palette: set[tuple[int, int, int]],
) -> tuple[bool, dict]:
    top = side_band_color(candidate_image, 0.10)
    middle = side_band_color(candidate_image, 0.50)
    bottom = side_band_color(candidate_image, 0.94)
    means = [statistics.mean(value) for value in (top, middle, bottom)]
    distances = [palette_distance(value, palette) for value in (top, bottom)]
    passed = (
        means[0] >= 230
        and means[2] >= 180
        and means[0] > means[1] > means[2]
        and means[0] - means[2] >= 10
        and bottom[2] >= bottom[0] + 5
        and max(bottom) - min(bottom) <= 65
        and max(distances) <= 50
    )
    return passed, {
        "top_rgb": list(top),
        "middle_rgb": list(middle),
        "bottom_rgb": list(bottom),
        "mean_brightness": [round(value, 3) for value in means],
        "picture_palette_distances": [round(value, 3) for value in distances],
    }


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {
        "criterion_id": identifier,
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
        "raw": {
            "description": description,
            "weight": weight,
            "type": kind,
            "passed": bool(passed),
        },
    }


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [
        item["criterion_id"]
        for item in criteria
        if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]
    ]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(
        item["raw"]["weight"]
        for item in criteria
        if item["raw"]["weight"] > 0 and item["raw"]["passed"]
    )
    penalty = sum(
        abs(item["raw"]["weight"])
        for item in criteria
        if item["raw"]["weight"] < 0 and not item["raw"]["passed"]
    )
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [sha256(path)] if path.is_file() else [],
        "criterion_results": criteria,
        "reward": reward,
        "passed": reward >= 0.8 and not hurdles,
        "hurdle_failures": hurdles,
    }


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        criteria = [
            result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配")
            for identifier in SPECS
        ]
        return make_result(path, criteria)
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        criteria = [
            result_item(identifier, False, f"PPTX 打开失败：{exc}")
            for identifier in SPECS
        ]
        return make_result(path, criteria)

    criteria: list[dict] = []
    package_passed, package_evidence = package_ok(path)
    size_passed = (
        len(source.slides) == 1
        and len(candidate.slides) == 1
        and candidate.slide_width == source.slide_width
        and candidate.slide_height == source.slide_height
    )
    criteria.append(
        result_item(
            "R001",
            package_passed and size_passed,
            f"{package_evidence}，输入页数={len(source.slides)}，"
            f"候选页数={len(candidate.slides)}，页面="
            f"{candidate.slide_width / EMU_PER_INCH:.4f}×"
            f"{candidate.slide_height / EMU_PER_INCH:.4f}英寸",
        )
    )

    try:
        accent_rgb = theme_accent1(source_path)
        background = background_inventory(path, accent_rgb)
    except Exception as exc:
        accent_rgb = (68, 114, 196)
        background = {"structure_passed": False, "color_passed": False, "error": str(exc)}
    criteria.append(
        result_item(
            "R002",
            bool(background.get("structure_passed")),
            f"原生背景渐变={json.dumps(background, ensure_ascii=False)}",
        )
    )

    try:
        palette = source_picture_palette(source)
        top_rgb = tuple(background.get("top_rgb") or (0, 0, 0))
        bottom_rgb = tuple(background.get("bottom_rgb") or (0, 0, 0))
        palette_distances = [
            palette_distance(top_rgb, palette),
            palette_distance(bottom_rgb, palette),
        ]
    except Exception:
        palette = set()
        palette_distances = [float("inf"), float("inf")]
    color_passed = bool(background.get("color_passed")) and max(palette_distances) <= 50
    criteria.append(
        result_item(
            "R003",
            color_passed,
            f"主题 accent1={accent_rgb}，结构颜色通过="
            f"{background.get('color_passed', False)}，与原图量化色集距离="
            f"{[round(value, 3) for value in palette_distances]}，"
            f"模式=theme:{background.get('theme_mode', False)}/"
            f"sampled-rgb:{background.get('sampled_rgb_mode', False)}",
        )
    )

    semantic_passed = (
        len(candidate.slides) == 1
        and slide_signature(candidate.slides[0]) == slide_signature(source.slides[0])
    )
    criteria.append(
        result_item(
            "R004",
            semantic_passed,
            f"页面语义对象签名保持={semantic_passed}，"
            f"对象数={len(source.slides[0].shapes)}→"
            f"{len(candidate.slides[0].shapes) if len(candidate.slides) else 0}",
        )
    )

    source_pictures = [
        shape for shape in source.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    candidate_pictures = [
        shape for shape in candidate.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ] if len(candidate.slides) else []
    picture_core_passed = False
    border_passed = False
    border = {"reason": "图片数不是 1"}
    if len(source_pictures) == 1 and len(candidate_pictures) == 1:
        source_picture = source_pictures[0]
        candidate_picture = candidate_pictures[0]
        picture_core_passed = (
            picture_blob_hash(candidate_picture) == picture_blob_hash(source_picture)
            and rounded_geometry(candidate_picture) == rounded_geometry(source_picture)
            and round(float(candidate_picture.crop_left), 6) == round(float(source_picture.crop_left), 6)
            and round(float(candidate_picture.crop_right), 6) == round(float(source_picture.crop_right), 6)
            and round(float(candidate_picture.crop_top), 6) == round(float(source_picture.crop_top), 6)
            and round(float(candidate_picture.crop_bottom), 6) == round(float(source_picture.crop_bottom), 6)
        )
        border_passed, border = border_evidence(candidate_picture, accent_rgb)
    criteria.append(
        result_item(
            "R005",
            picture_core_passed and border_passed,
            f"图片字节/裁剪/几何保持={picture_core_passed}，"
            f"边框={json.dumps(border, ensure_ascii=False)}",
        )
    )

    try:
        source_media = media_hashes(source_path)
        candidate_media = media_hashes(path)
        media_equal = source_media == candidate_media
        source_members = stable_member_names(source_path)
        candidate_members = stable_member_names(path)
        member_sets_equal = source_members == candidate_members
    except Exception:
        source_media = Counter()
        candidate_media = Counter()
        media_equal = False
        member_sets_equal = False
        source_members = set()
        candidate_members = set()
    flat_pictures = full_slide_pictures(candidate)
    native_structure_passed = (
        package_passed
        and len(candidate.slides) == 1
        and len(candidate.slides[0].shapes) == 3
        and len(candidate_pictures) == 1
        and media_equal
        and member_sets_equal
        and not flat_pictures
        and bool(background.get("structure_passed"))
    )
    criteria.append(
        result_item(
            "R008",
            native_structure_passed,
            f"原生对象数={len(candidate.slides[0].shapes) if len(candidate.slides) else 0}，"
            f"媒体哈希相等={media_equal}，媒体数="
            f"{sum(source_media.values())}→{sum(candidate_media.values())}，"
            f"包成员集相等={member_sets_equal}，整页图片={flat_pictures}",
        )
    )

    artifact_files = sorted(
        str(item.relative_to(artifact_dir))
        for item in artifact_dir.rglob("*")
        if item.is_file()
    )
    all_text = "\n".join(
        shape.text
        for slide in candidate.slides
        for shape in iter_shapes(slide.shapes)
        if getattr(shape, "has_text_frame", False)
    )
    lowered = all_text.lower()
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker.lower() in lowered]
    hygiene_passed = artifact_files == [OUTPUT_NAME] and not markers
    criteria.append(
        result_item(
            "R009",
            hygiene_passed,
            f"输出文件={artifact_files}，可疑标记={markers}",
        )
    )

    try:
        source_content = normalized_content_hash(source_path, remove_picture_line=False)
        candidate_content = normalized_content_hash(path, remove_picture_line=border_passed)
        exact_content = source_content == candidate_content
        critical_equal = critical_hashes(source_path) == critical_hashes(path)
    except Exception:
        exact_content = False
        critical_equal = False
    criteria.append(
        result_item(
            "R010",
            exact_content and critical_equal,
            f"去除背景/合法边框后 cSld 保持={exact_content}，"
            f"关键非幻灯片包部件哈希保持={critical_equal}",
        )
    )

    pre_render_ready = (
        package_passed
        and size_passed
        and bool(background.get("structure_passed"))
        and color_passed
        and semantic_passed
        and picture_core_passed
        and border_passed
        and native_structure_passed
    )
    picture_render_passed = False
    picture_similarity = 0.0
    rendered_background_passed = False
    rendered_background = {"reason": "前置检查失败"}
    render_evidence = "前置结构、颜色或保持检查失败，未执行渲染"
    if pre_render_ready:
        with tempfile.TemporaryDirectory(prefix="pptarena041-judge-") as temporary:
            root = Path(temporary)
            source_images, source_pages, source_note = render_presentation(
                source_path, root / "source"
            )
            candidate_images, candidate_pages, candidate_note = render_presentation(
                path, root / "candidate"
            )
            render_evidence = f"源[{source_note}]；候选[{candidate_note}]"
            if len(source_images) == 1 and len(candidate_images) == 1:
                with Image.open(source_images[0]) as source_image, Image.open(candidate_images[0]) as candidate_image:
                    source_crop = picture_crop(
                        source_image,
                        source_pictures[0],
                        (source.slide_width, source.slide_height),
                    )
                    candidate_crop = picture_crop(
                        candidate_image,
                        candidate_pictures[0],
                        (candidate.slide_width, candidate.slide_height),
                    )
                    picture_similarity = image_similarity(source_crop, candidate_crop)
                    rendered_background_passed, rendered_background = rendered_background_evidence(
                        candidate_image, palette
                    )
            picture_render_passed = (
                source_pages == 1
                and candidate_pages == 1
                and len(source_images) == 1
                and len(candidate_images) == 1
                and picture_similarity >= 0.9985
            )
    criteria.append(
        result_item(
            "R006",
            picture_render_passed,
            f"中央图片内部像素相似度={picture_similarity:.6f}；{render_evidence}",
        )
    )
    criteria.append(
        result_item(
            "R007",
            rendered_background_passed,
            f"实际背景={json.dumps(rendered_background, ensure_ascii=False)}；{render_evidence}",
        )
    )

    severe = (
        not pre_render_ready
        or not picture_render_passed
        or not rendered_background_passed
        or bool(flat_pictures)
        or bool(markers)
    )
    criteria.append(result_item("R011", not severe, f"严重错误={severe}"))
    order = {identifier: index for index, identifier in enumerate(SPECS)}
    criteria.sort(key=lambda item: order[item["criterion_id"]])
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    output = args.artifact_dir / OUTPUT_NAME
    if output.is_file():
        evaluation = evaluate(output, args.artifact_dir)
    else:
        evaluation = {
            "task_id": TASK_ID,
            "rubric_version": RUBRIC_VERSION,
            "judge_version": JUDGE_VERSION,
            "artifact_sha256": [],
            "criterion_results": [],
            "reward": 0.0,
            "passed": False,
            "hurdle_failures": ["missing_output"],
        }
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(
        json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
