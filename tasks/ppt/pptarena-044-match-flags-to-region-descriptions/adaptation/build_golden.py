#!/usr/bin/env python3
"""PPTArena Case 44 专用 Golden 构建器。"""

from __future__ import annotations

import argparse
import hashlib
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt


INPUT_SHA256 = "8c2c6f409684507f2d56cc205c8a6b17d69e65d5036748dbe6d2098e24f325e2"
TITLE = "African Dialogue & U.S. FTC"
SOUTH_AFRICA_HASH = "4f1abe2384af08da1d2113c8395095411f0e8b86cfb7db9f154d1137844af5b0"
PAIR_TEXTS = {
    "f6ad8649736281f1d3d1d6093e8f2654251a0a8b061cd4ced45c895374bc0188":
        "EAST AFRICA - Kenyan Competition Authority",
    "021690c3c58a55242b53d43e07bb0c0daf4d0f1bd70aabb12a2feb1c7757bb1a":
        "WEST AFRICA - Nigerian Consumer Protection Council",
    "8b98a3cf1c84aa7aac4a45637b7da811f38ea7d1ba2ea1d0c8f27310b107d80e":
        "NORTH AFRICA - Egyptian Consumer Protection Agency",
    "128d4315add1153b7be16b728dc578f5d5f955f9ac245e428fda0f2adac4c8a8":
        "SOUTHERN AFRICA - Zambian Competition and Consumer Protection Commission",
}
GRID_PLACEMENTS = {
    list(PAIR_TEXTS)[0]: (0.55, 2.25, 2.55, 2.50, 1.35),
    list(PAIR_TEXTS)[1]: (5.15, 2.25, 7.15, 2.50, 1.35),
    list(PAIR_TEXTS)[2]: (0.55, 4.85, 2.55, 5.10, 1.35),
    list(PAIR_TEXTS)[3]: (5.15, 4.85, 7.15, 5.00, 1.60),
}
ROW_PLACEMENTS = {
    image_hash: (0.70, 2.05 + index * 1.25, 2.20, 2.22 + index * 1.25, 0.75)
    for index, image_hash in enumerate(PAIR_TEXTS)
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalize(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def picture_hash(shape) -> str:
    return hashlib.sha256(shape.image.blob).hexdigest()


def remove_shape(shape) -> None:
    shape._element.getparent().remove(shape._element)


def add_heading(slide) -> None:
    box = slide.shapes.add_textbox(Inches(0.55), Inches(1.52), Inches(8.90), Inches(0.55))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = "Regional Liaisons"
    run.font.name = "Arial"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(63, 61, 38)


def add_label(slide, text: str, left: float, top: float, width: float, height: float, font_size: float) -> None:
    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.05)
    frame.margin_right = Inches(0.02)
    frame.margin_top = frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.line_spacing = 1.05
    prefix = text.split(" - ", 1)[0] + " - "
    region = paragraph.add_run()
    region.text = prefix
    region.font.name = "Arial"
    region.font.size = Pt(font_size)
    region.font.bold = True
    region.font.color.rgb = RGBColor(35, 35, 35)
    organization = paragraph.add_run()
    organization.text = text[len(prefix):]
    organization.font.name = "Arial"
    organization.font.size = Pt(font_size)
    organization.font.color.rgb = RGBColor(35, 35, 35)


def build(source: Path, output: Path, layout: str = "grid") -> None:
    if sha256(source) != INPUT_SHA256:
        raise ValueError("冻结输入 SHA256 不匹配")
    if layout not in {"grid", "rows"}:
        raise ValueError(f"不支持的本题排版：{layout}")

    presentation = Presentation(source)
    if len(presentation.slides) != 1:
        raise ValueError("冻结输入不再是 1 页")
    slide = presentation.slides[0]
    if normalize(slide.shapes.title.text if slide.shapes.title else "") != TITLE:
        raise ValueError("冻结输入标题不匹配")

    content_placeholders = [shape for shape in slide.shapes if shape.shape_id == 3]
    if len(content_placeholders) != 1:
        raise ValueError("无法唯一定位原项目符号列表")

    flags = {}
    extra_flags = []
    for shape in list(slide.shapes):
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        image_hash = picture_hash(shape)
        if image_hash in PAIR_TEXTS:
            flags[image_hash] = shape
        elif image_hash == SOUTH_AFRICA_HASH:
            extra_flags.append(shape)
    if set(flags) != set(PAIR_TEXTS) or len(extra_flags) != 1:
        raise ValueError("冻结输入中的 4 面正确国旗或隐藏南非国旗异常")

    remove_shape(content_placeholders[0])
    remove_shape(extra_flags[0])
    add_heading(slide)

    placements = GRID_PLACEMENTS if layout == "grid" else ROW_PLACEMENTS
    for image_hash, text in PAIR_TEXTS.items():
        flag_left, flag_top, label_left, label_top, label_height = placements[image_hash]
        flag = flags[image_hash]
        flag.left = Inches(flag_left)
        flag.top = Inches(flag_top)
        if layout == "grid":
            flag.width, flag.height = Inches(1.75), Inches(1.40)
            label_width, font_size = 2.35, 15
        else:
            flag.width, flag.height = Inches(1.20), Inches(0.96)
            label_width, font_size = 7.15, 17
        add_label(
            slide,
            text,
            label_left,
            label_top,
            label_width,
            label_height,
            font_size,
        )

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    validate_output(output)


def validate_output(path: Path) -> None:
    presentation = Presentation(path)
    if len(presentation.slides) != 1:
        raise ValueError("产物页数不是 1")
    slide = presentation.slides[0]
    if normalize(slide.shapes.title.text if slide.shapes.title else "") != TITLE:
        raise ValueError("产物标题未保留")
    text = "\n".join(
        shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)
    )
    if any(normalize(value) not in normalize(text) for value in PAIR_TEXTS.values()):
        raise ValueError("产物缺少地区机构说明")
    hashes = [
        picture_hash(shape)
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    if any(hashes.count(image_hash) != 1 for image_hash in PAIR_TEXTS):
        raise ValueError("四面正确国旗未各保留一次")
    if SOUTH_AFRICA_HASH in hashes:
        raise ValueError("隐藏的南非国旗未删除")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--layout", choices=["grid", "rows"], default="grid")
    args = parser.parse_args()
    build(args.source, args.output, layout=args.layout)
    print(f"{args.output}\nsha256={sha256(args.output)}")


if __name__ == "__main__":
    main()
