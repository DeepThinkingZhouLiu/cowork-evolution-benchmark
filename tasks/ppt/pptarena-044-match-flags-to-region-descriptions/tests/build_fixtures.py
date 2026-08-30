#!/usr/bin/env python3
"""只为 PPTArena Case 44 构造合法变体、负例和中等例。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import shutil
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt


OUTPUT_NAME = "MatchFlagtoText_GroundTruthA.pptx"


def load_builder(path: Path):
    spec = importlib.util.spec_from_file_location("pptarena044_builder", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("无法加载本题 Golden 构建器")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def output_path(root: Path, case: str) -> Path:
    target = root / case / OUTPUT_NAME
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def image_hash(shape) -> str:
    return hashlib.sha256(shape.image.blob).hexdigest()


def find_picture(presentation: Presentation, expected_hash: str):
    matches = [
        shape
        for shape in presentation.slides[0].shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and image_hash(shape) == expected_hash
    ]
    if len(matches) != 1:
        raise ValueError(f"无法唯一定位图片 {expected_hash[:16]}")
    return matches[0]


def find_label(presentation: Presentation, text: str):
    matches = [
        shape
        for shape in presentation.slides[0].shapes
        if getattr(shape, "has_text_frame", False) and " ".join(shape.text.split()) == text
    ]
    if len(matches) != 1:
        raise ValueError(f"无法唯一定位说明：{text}")
    return matches[0]


def remove_shape(shape) -> None:
    shape._element.getparent().remove(shape._element)


def save_copy(golden: Path, output: Path, mutator) -> None:
    presentation = Presentation(golden)
    mutator(presentation)
    presentation.save(output)


def valid_alternate(source: Path, builder, output: Path) -> None:
    builder.build(source, output, layout="rows")


def inconsistent_flag_size(golden: Path, builder, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        flag = find_picture(presentation, list(builder.PAIR_TEXTS)[0])
        flag.width = Inches(1.25)
        flag.height = Inches(1.00)
        label = find_label(presentation, list(builder.PAIR_TEXTS.values())[0])
        label.left = Inches(2.05)

    save_copy(golden, output, mutate)


def add_south_africa(golden: Path, source: Path, builder, output: Path) -> None:
    source_presentation = Presentation(source)
    extra = find_picture(source_presentation, builder.SOUTH_AFRICA_HASH)

    def mutate(presentation: Presentation) -> None:
        presentation.slides[0].shapes.add_picture(
            BytesIO(extra.image.blob), Inches(7.45), Inches(5.85), Inches(1.25), Inches(0.90)
        )

    save_copy(golden, output, mutate)


def missing_flag(golden: Path, builder, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        remove_shape(find_picture(presentation, list(builder.PAIR_TEXTS)[2]))

    save_copy(golden, output, mutate)


def duplicate_flag(golden: Path, builder, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        flag = find_picture(presentation, list(builder.PAIR_TEXTS)[0])
        presentation.slides[0].shapes.add_picture(
            BytesIO(flag.image.blob), Inches(3.85), Inches(6.10), Inches(1.10), Inches(0.88)
        )

    save_copy(golden, output, mutate)


def swapped_labels(golden: Path, builder, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        first = find_label(presentation, list(builder.PAIR_TEXTS.values())[0])
        second = find_label(presentation, list(builder.PAIR_TEXTS.values())[1])
        first_geometry = first.left, first.top, first.width, first.height
        second_geometry = second.left, second.top, second.width, second.height
        first.left, first.top, first.width, first.height = second_geometry
        second.left, second.top, second.width, second.height = first_geometry

    save_copy(golden, output, mutate)


def truncated_label(golden: Path, builder, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        label = find_label(presentation, list(builder.PAIR_TEXTS.values())[3])
        label.text = "SOUTHERN AFRICA - Zambian Competition Commission"

    save_copy(golden, output, mutate)


def combined_list(golden: Path, builder, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        slide = presentation.slides[0]
        for text in builder.PAIR_TEXTS.values():
            remove_shape(find_label(presentation, text))
        box = slide.shapes.add_textbox(Inches(2.45), Inches(2.10), Inches(7.10), Inches(4.70))
        box.text_frame.text = "\n\n".join(builder.PAIR_TEXTS.values())
        for paragraph in box.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(16)

    save_copy(golden, output, mutate)


def changed_title(golden: Path, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        presentation.slides[0].shapes.title.text = "African Dialogue"

    save_copy(golden, output, mutate)


def removed_header_image(golden: Path, builder, output: Path) -> None:
    # 顶图哈希不在 PAIR_TEXTS，直接按页面左上图定位。
    def actual_mutate(presentation: Presentation) -> None:
        pictures = [shape for shape in presentation.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        header = min(pictures, key=lambda shape: shape.top)
        remove_shape(header)

    save_copy(golden, output, actual_mutate)


def moved_header_image(golden: Path, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        pictures = [shape for shape in presentation.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        header = min(pictures, key=lambda shape: shape.top)
        header.left += Inches(0.30)

    save_copy(golden, output, mutate)


def offslide_label(golden: Path, builder, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        find_label(presentation, list(builder.PAIR_TEXTS.values())[0]).left = Inches(-0.50)

    save_copy(golden, output, mutate)


def overlapping_pairs(golden: Path, builder, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        hashes = list(builder.PAIR_TEXTS)
        texts = list(builder.PAIR_TEXTS.values())
        target_flag = find_picture(presentation, hashes[0])
        moved_flag = find_picture(presentation, hashes[1])
        moved_flag.left, moved_flag.top = target_flag.left, target_flag.top
        target_label = find_label(presentation, texts[0])
        moved_label = find_label(presentation, texts[1])
        moved_label.left, moved_label.top = target_label.left, target_label.top

    save_copy(golden, output, mutate)


def flattened(golden: Path, slide_image: Path, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        slide = presentation.slides[0]
        for shape in list(slide.shapes):
            remove_shape(shape)
        slide.shapes.add_picture(
            str(slide_image), 0, 0,
            width=presentation.slide_width, height=presentation.slide_height,
        )

    save_copy(golden, output, mutate)


def hardcoded_marker(golden: Path, output: Path) -> None:
    def mutate(presentation: Presentation) -> None:
        headings = [shape for shape in presentation.slides[0].shapes if getattr(shape, "has_text_frame", False) and shape.text.strip() == "Regional Liaisons"]
        if len(headings) != 1:
            raise ValueError("无法定位小标题")
        headings[0].text = "Regional Liaisons\nsolution/Golden"

    save_copy(golden, output, mutate)


def extra_slide(golden: Path, output: Path) -> None:
    save_copy(golden, output, lambda presentation: presentation.slides.add_slide(presentation.slide_layouts[0]))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--builder", type=Path, required=True)
    parser.add_argument("--slide-image", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()
    builder = load_builder(args.builder)

    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))
    valid_alternate(args.source, builder, output_path(args.outdir, "valid-alternate"))
    inconsistent_flag_size(args.golden, builder, output_path(args.outdir, "inconsistent-flag-size"))
    add_south_africa(args.golden, args.source, builder, output_path(args.outdir, "extra-south-africa"))
    missing_flag(args.golden, builder, output_path(args.outdir, "missing-flag"))
    duplicate_flag(args.golden, builder, output_path(args.outdir, "duplicate-flag"))
    swapped_labels(args.golden, builder, output_path(args.outdir, "swapped-labels"))
    truncated_label(args.golden, builder, output_path(args.outdir, "truncated-label"))
    combined_list(args.golden, builder, output_path(args.outdir, "combined-list"))
    changed_title(args.golden, output_path(args.outdir, "changed-title"))
    removed_header_image(args.golden, builder, output_path(args.outdir, "removed-header-image"))
    moved_header_image(args.golden, output_path(args.outdir, "moved-header-image"))
    offslide_label(args.golden, builder, output_path(args.outdir, "offslide-label"))
    overlapping_pairs(args.golden, builder, output_path(args.outdir, "overlapping-pairs"))
    flattened(args.golden, args.slide_image, output_path(args.outdir, "flattened"))
    hardcoded_marker(args.golden, output_path(args.outdir, "hardcoded-marker"))
    extra_slide(args.golden, output_path(args.outdir, "extra-slide"))


if __name__ == "__main__":
    main()
