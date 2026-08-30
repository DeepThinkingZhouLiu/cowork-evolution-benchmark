#!/usr/bin/env python3
"""PPTArena Case 44 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import posixpath
import re
import statistics
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from PIL import Image, ImageChops, ImageStat
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "MatchFlagtoText_GroundTruthA.pptx"
INPUT_NAME = "MatchFlagtoText_TestA.pptx"
INPUT_SHA256 = "8c2c6f409684507f2d56cc205c8a6b17d69e65d5036748dbe6d2098e24f325e2"
TASK_ID = "cowork-evo/pptarena-044-match-flags-to-region-descriptions"
RUBRIC_VERSION = "pptarena-case-044-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-044-deterministic-v0.1"
EMU_PER_INCH = 914_400
TITLE = "African Dialogue & U.S. FTC"
HEADER_HASHES = {
    "8fdf0cdb55f4a69ecaf6b049fd2679be63e23d017eea2806000dee51a076bf03",
    "f8aa0bc7a2b33f96ebd7b88659a0316a0c43750f0b5749b61adf78e4368a3f20",
}
SOUTH_AFRICA_HASH = "4f1abe2384af08da1d2113c8395095411f0e8b86cfb7db9f154d1137844af5b0"
PAIR_TEXTS = {
    "f6ad8649736281f1d3d1d6093e8f2654251a0a8b061cd4ced45c895374bc0188":
        "EAST AFRICA - Kenyan Competition Authority",
    "021690c3c58a55242b53d43e07bb0c0daf4d0f1bd70aabb12a2feb1c7757bb1a":
        "WEST AFRICA - Nigerian Consumer Protection Council",
    "8b98a3cf1c84aa7aac4a45637b7da811f38ea7d1ba2ea1d0c8f27310b107d80e":
        "NORTH AFRICA - Egyptian Consumer Protection Agency",
    "128d4315add1153b7be16b728dc578f5d5f955f9ac245e428fda0f2adac4c8a8":
        "SOUTHERN AFRICA - Zambian Competition and Consumer Protection Commission",
}
SUSPICIOUS_MARKERS = ["solution/golden", "judge-result", "artifact_dir", "pptarena-case-044", "/root/"]

SPECS = {
    "R001": ("PPTX 完整且仍为原尺寸单页文稿", 7, "hurdle"),
    "R002": ("四面正确国旗各保留一次且南非国旗已删", 18, "hurdle"),
    "R003": ("四条完整说明为独立文字对象且原列表已删", 16, "hurdle"),
    "R004": ("四条说明与最近国旗的国家映射正确且清楚", 20, "hurdle"),
    "R005": ("主标题、顶部两图与备注保持", 12, "hurdle"),
    "R006": ("配对对象在页内且不重叠，可编辑且无整页截图", 10, "hurdle"),
    "R007": ("国旗尺寸和配对间距稳定，排版清楚", 7, "positive"),
    "R008": ("LibreOffice 渲染中标题、顶图和四条说明可见", 6, "hurdle"),
    "R009": ("PPTX 关系和所需媒体完整，无宏", 2, "hurdle"),
    "R010": ("输出目录干净且无答案标记", 2, "positive"),
    "R011": ("无错配、额外国旗、标题顶图丢失、越界重叠或可编辑性丧失", -20, "penalty"),
}


def normalize(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def locate_source() -> Path | None:
    candidates = [
        Path("/app/input_files") / INPUT_NAME,
        Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME,
    ]
    return next((path for path in candidates if path.is_file()), None)


def picture_hash(shape) -> str:
    return hashlib.sha256(shape.image.blob).hexdigest()


def geometry(shape) -> tuple[int, int, int, int]:
    return int(shape.left), int(shape.top), int(shape.width), int(shape.height)


def safe_value(callback, default=""):
    try:
        return callback()
    except Exception:
        return default


def color_signature(color_format) -> tuple:
    return tuple(
        safe_value(lambda name=name: str(getattr(color_format, name)))
        for name in ["type", "rgb", "theme_color", "brightness"]
    )


def font_signature(font) -> tuple:
    return (
        font.name,
        round(font.size.pt, 4) if font.size else None,
        font.bold,
        font.italic,
        font.underline,
        safe_value(lambda: str(font.language_id)),
        color_signature(font.color),
    )


def text_signature(shape) -> tuple:
    frame = shape.text_frame
    paragraphs = []
    for paragraph in frame.paragraphs:
        paragraphs.append(
            (
                paragraph.level,
                str(paragraph.alignment),
                str(paragraph.line_spacing),
                str(paragraph.space_before),
                str(paragraph.space_after),
                tuple((run.text, font_signature(run.font)) for run in paragraph.runs),
            )
        )
    return (
        normalize(shape.text),
        tuple(round(value / 1000) * 1000 for value in geometry(shape)),
        (
            frame.margin_left,
            frame.margin_right,
            frame.margin_top,
            frame.margin_bottom,
            str(frame.vertical_anchor),
            frame.word_wrap,
            str(frame.auto_size),
        ),
        tuple(paragraphs),
    )


def title_signature(shape) -> tuple:
    return (
        str(shape.shape_type),
        str(shape.placeholder_format.type) if shape.is_placeholder else "",
        text_signature(shape),
    )


def rectangle_distance(first, second) -> float:
    left1, top1, width1, height1 = geometry(first)
    left2, top2, width2, height2 = geometry(second)
    horizontal = max(left1 - (left2 + width2), left2 - (left1 + width1), 0)
    vertical = max(top1 - (top2 + height2), top2 - (top1 + height1), 0)
    return math.hypot(horizontal, vertical) / EMU_PER_INCH


def overlap_area(first, second) -> float:
    left1, top1, width1, height1 = geometry(first)
    left2, top2, width2, height2 = geometry(second)
    overlap_width = max(0, min(left1 + width1, left2 + width2) - max(left1, left2))
    overlap_height = max(0, min(top1 + height1, top2 + height2) - max(top1, top2))
    return overlap_width * overlap_height / (EMU_PER_INCH**2)


def inside_slide(shape, slide_width: int, slide_height: int, tolerance_in: float = 0.03) -> bool:
    tolerance = int(tolerance_in * EMU_PER_INCH)
    left, top, width, height = geometry(shape)
    return (
        left >= -tolerance
        and top >= -tolerance
        and left + width <= slide_width + tolerance
        and top + height <= slide_height + tolerance
    )


def package_checks(path: Path) -> tuple[bool, dict]:
    evidence = {"broken_member": None, "missing_targets": [], "macros": []}
    try:
        with zipfile.ZipFile(path) as archive:
            evidence["broken_member"] = archive.testzip()
            names = set(archive.namelist())
            evidence["macros"] = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
            for rels_name in sorted(name for name in names if name.endswith(".rels")):
                root = __import__("xml.etree.ElementTree", fromlist=["ElementTree"]).fromstring(
                    archive.read(rels_name)
                )
                if rels_name == "_rels/.rels":
                    owner = ""
                else:
                    prefix, filename = rels_name.split("/_rels/", 1)
                    owner = posixpath.join(prefix, filename[:-5])
                base = posixpath.dirname(owner)
                for relation in root:
                    if relation.attrib.get("TargetMode") == "External":
                        continue
                    target = relation.attrib.get("Target", "")
                    resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                    if resolved not in names:
                        evidence["missing_targets"].append(f"{rels_name} -> {resolved}")
        passed = evidence["broken_member"] is None and not evidence["missing_targets"] and not evidence["macros"]
        return passed, evidence
    except Exception as exc:
        evidence["exception"] = str(exc)
        return False, evidence


def render_one(path: Path, outdir: Path) -> tuple[Path | None, Path | None, int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "lo-profile"
    profile.mkdir()
    completed = subprocess.run(
        [
            "soffice",
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(outdir),
            str(path),
        ],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return None, None, 0, completed.stdout[-500:]
    info = subprocess.run(
        ["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
        text=True, timeout=30, check=False
    ).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    pages = int(match.group(1)) if match else 0
    prefix = outdir / "slide"
    raster = subprocess.run(
        ["pdftoppm", "-singlefile", "-png", "-r", "120", str(pdf), str(prefix)],
        stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=240, check=False
    )
    image = outdir / "slide.png"
    return pdf, image if raster.returncode == 0 and image.is_file() else None, pages, f"soffice={completed.returncode}, raster={raster.returncode}"


def extract_pdf_text(pdf: Path) -> str:
    result = subprocess.run(
        ["pdftotext", "-layout", str(pdf), "-"],
        stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=60, check=False
    )
    return normalize(result.stdout) if result.returncode == 0 else ""


def top_similarity(source_image: Path, candidate_image: Path) -> float:
    with Image.open(source_image) as first, Image.open(candidate_image) as second:
        left, right = first.convert("RGB"), second.convert("RGB")
        if left.size != right.size:
            return 0.0
        crop_height = int(left.height * 0.195)
        difference = ImageChops.difference(
            left.crop((0, 0, left.width, crop_height)),
            right.crop((0, 0, right.width, crop_height)),
        )
        rms = statistics.mean(ImageStat.Stat(difference).rms)
        return max(0.0, 1.0 - rms / 255.0)


def media_hashes(path: Path) -> Counter[str]:
    with zipfile.ZipFile(path) as archive:
        return Counter(
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("ppt/media/")
        )


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {
        "criterion_id": identifier,
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
        "raw": {"description": description, "weight": weight, "type": kind, "passed": bool(passed)},
    }


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [item["criterion_id"] for item in criteria if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0 and item["raw"]["passed"])
    penalty = sum(abs(item["raw"]["weight"]) for item in criteria if item["raw"]["weight"] < 0 and not item["raw"]["passed"])
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [sha256(path)] if path.is_file() else [],
        "criterion_results": criteria,
        "reward": reward,
        "passed": reward >= 0.8 and not hurdles,
        "hurdle_failures": hurdles,
    }


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        return make_result(path, [result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配") for identifier in SPECS])
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        return make_result(path, [result_item(identifier, False, f"PPTX 打开失败：{exc}") for identifier in SPECS])

    criteria: list[dict] = []
    package_passed, package_evidence = package_checks(path)
    size_passed = (
        len(source.slides) == 1 and len(candidate.slides) == 1
        and candidate.slide_width == source.slide_width
        and candidate.slide_height == source.slide_height
    )
    criteria.append(result_item("R001", package_passed and size_passed, f"页数={len(candidate.slides)}，尺寸一致={size_passed}，{json.dumps(package_evidence, ensure_ascii=False)}"))

    slide = candidate.slides[0]
    pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    picture_map: dict[str, list] = {}
    for shape in pictures:
        picture_map.setdefault(picture_hash(shape), []).append(shape)
    flags_passed = all(len(picture_map.get(image_hash, [])) == 1 for image_hash in PAIR_TEXTS)
    south_africa_absent = not picture_map.get(SOUTH_AFRICA_HASH)
    criteria.append(result_item("R002", flags_passed and south_africa_absent, f"正确国旗计数={{{', '.join(f'{key[:8]}:{len(picture_map.get(key, []))}' for key in PAIR_TEXTS)}}}，南非国旗数={len(picture_map.get(SOUTH_AFRICA_HASH, []))}"))

    text_shapes = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False) and normalize(shape.text)]
    label_map = {}
    label_details = []
    descriptions_passed = True
    for image_hash, expected in PAIR_TEXTS.items():
        matches = [shape for shape in text_shapes if normalize(shape.text) == expected]
        label_details.append({"text": expected, "matches": len(matches)})
        if len(matches) != 1:
            descriptions_passed = False
        else:
            label_map[image_hash] = matches[0]
    combined_list = any(sum(expected in normalize(shape.text) for expected in PAIR_TEXTS.values()) > 1 for shape in text_shapes)
    criteria.append(result_item("R003", descriptions_passed and not combined_list, f"独立说明={json.dumps(label_details, ensure_ascii=False)}，合并列表={combined_list}"))

    pair_details = []
    pairing_passed = flags_passed and descriptions_passed
    pair_gaps = []
    if pairing_passed:
        all_flags = {image_hash: picture_map[image_hash][0] for image_hash in PAIR_TEXTS}
        for image_hash, label in label_map.items():
            distances = {flag_hash: rectangle_distance(label, flag) for flag_hash, flag in all_flags.items()}
            nearest = min(distances, key=distances.get)
            correct_distance = distances[image_hash]
            passed = nearest == image_hash and correct_distance <= 0.80
            pair_gaps.append(correct_distance)
            pair_details.append({"label": PAIR_TEXTS[image_hash], "nearest_hash": nearest[:16], "correct_hash": image_hash[:16], "gap_in": round(correct_distance, 4), "passed": passed})
            pairing_passed = pairing_passed and passed
    criteria.append(result_item("R004", pairing_passed, json.dumps(pair_details, ensure_ascii=False)))

    source_slide = source.slides[0]
    source_title = next((shape for shape in source_slide.shapes if getattr(shape, "has_text_frame", False) and normalize(shape.text) == TITLE), None)
    candidate_titles = [shape for shape in text_shapes if normalize(shape.text) == TITLE]
    title_passed = source_title is not None and len(candidate_titles) == 1 and title_signature(candidate_titles[0]) == title_signature(source_title)
    header_details = []
    header_passed = True
    source_headers = {picture_hash(shape): shape for shape in source_slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and picture_hash(shape) in HEADER_HASHES}
    for image_hash in HEADER_HASHES:
        matches = picture_map.get(image_hash, [])
        passed = len(matches) == 1 and image_hash in source_headers and geometry(matches[0]) == geometry(source_headers[image_hash])
        header_details.append({"hash": image_hash[:16], "matches": len(matches), "geometry_preserved": passed})
        header_passed = header_passed and passed
    notes_passed = normalize(slide.notes_slide.notes_text_frame.text) == normalize(source_slide.notes_slide.notes_text_frame.text)
    criteria.append(result_item("R005", title_passed and header_passed and notes_passed, f"标题保持={title_passed}，顶图={header_details}，备注保持={notes_passed}"))

    pair_objects = []
    if flags_passed and descriptions_passed:
        pair_objects = [picture_map[image_hash][0] for image_hash in PAIR_TEXTS] + [label_map[image_hash] for image_hash in PAIR_TEXTS]
    inside = bool(pair_objects) and all(inside_slide(shape, candidate.slide_width, candidate.slide_height) for shape in pair_objects)
    overlaps = []
    for first_index in range(len(pair_objects)):
        for second_index in range(first_index + 1, len(pair_objects)):
            area = overlap_area(pair_objects[first_index], pair_objects[second_index])
            if area > 0.005:
                overlaps.append([first_index, second_index, round(area, 4)])
    page_area = candidate.slide_width * candidate.slide_height
    full_slide_pictures = [shape.name for shape in pictures if shape.width * shape.height / page_area >= 0.85]
    editable_passed = bool(pair_objects) and inside and not overlaps and not full_slide_pictures
    criteria.append(result_item("R006", editable_passed, f"配对对象数={len(pair_objects)}，全在页内={inside}，重叠={overlaps}，整页图={full_slide_pictures}"))

    flag_areas = []
    if flags_passed:
        flag_areas = [picture_map[image_hash][0].width * picture_map[image_hash][0].height / (EMU_PER_INCH**2) for image_hash in PAIR_TEXTS]
    size_ratio = max(flag_areas) / min(flag_areas) if flag_areas and min(flag_areas) > 0 else float("inf")
    gap_std = statistics.pstdev(pair_gaps) if len(pair_gaps) == 4 else float("inf")
    layout_passed = (
        editable_passed and size_ratio <= 1.25 and gap_std <= 0.30
        and all(0.70 <= area <= 4.50 for area in flag_areas)
    )
    criteria.append(result_item("R007", layout_passed, f"国旗面积={list(map(lambda value: round(value, 4), flag_areas))}，最大最小比={size_ratio:.4f}，配对间距标准差={gap_std:.4f}"))

    pre_render_ready = all(next(item for item in criteria if item["criterion_id"] == identifier)["raw"]["passed"] for identifier in ["R001", "R002", "R003", "R004", "R005", "R006"])
    render_passed = False
    render_evidence = {"skipped": "前置检查失败"}
    if pre_render_ready:
        with tempfile.TemporaryDirectory(prefix="pptarena044-judge-") as temporary:
            root = Path(temporary)
            source_pdf, source_image, source_pages, source_note = render_one(source_path, root / "source")
            candidate_pdf, candidate_image, candidate_pages, candidate_note = render_one(path, root / "candidate")
            pdf_text = extract_pdf_text(candidate_pdf) if candidate_pdf else ""
            anchors = [
                "African", "Dialogue", "FTC", "Kenyan", "Authority", "Nigerian",
                "Council", "Egyptian", "Agency", "Zambian", "Commission",
            ]
            anchors_found = [anchor for anchor in anchors if anchor.lower() in pdf_text.lower()]
            similarity = top_similarity(source_image, candidate_image) if source_image and candidate_image else 0.0
            label_dark_ratios = []
            if candidate_image:
                with Image.open(candidate_image) as rendered:
                    gray = rendered.convert("L")
                    for image_hash in PAIR_TEXTS:
                        label = label_map[image_hash]
                        left, top, width, height = geometry(label)
                        crop = gray.crop((
                            max(0, int(left * gray.width / candidate.slide_width)),
                            max(0, int(top * gray.height / candidate.slide_height)),
                            min(gray.width, int(math.ceil((left + width) * gray.width / candidate.slide_width))),
                            min(gray.height, int(math.ceil((top + height) * gray.height / candidate.slide_height))),
                        ))
                        histogram = crop.histogram()
                        dark_ratio = sum(histogram[:210]) / max(1, crop.width * crop.height)
                        label_dark_ratios.append(round(dark_ratio, 6))
            render_passed = (
                source_pages == 1
                and candidate_pages == 1
                and len(anchors_found) == len(anchors)
                and len(label_dark_ratios) == 4
                and all(value >= 0.005 for value in label_dark_ratios)
                and similarity >= 0.995
            )
            render_evidence = {
                "source": source_note,
                "candidate": candidate_note,
                "source_pages": source_pages,
                "candidate_pages": candidate_pages,
                "anchors_found": anchors_found,
                "expected_anchors": len(anchors),
                "label_dark_ratios": label_dark_ratios,
                "top_similarity": round(similarity, 6),
            }
    criteria.append(result_item("R008", render_passed, json.dumps(render_evidence, ensure_ascii=False)))

    candidate_media = media_hashes(path)
    expected_media = HEADER_HASHES | set(PAIR_TEXTS)
    media_passed = all(candidate_media[image_hash] >= 1 for image_hash in expected_media)
    criteria.append(result_item("R009", package_passed and media_passed, f"所需 6 个媒体保留={media_passed}，包={json.dumps(package_evidence, ensure_ascii=False)}"))

    artifact_files = sorted(str(item.relative_to(artifact_dir)) for item in artifact_dir.rglob("*") if item.is_file())
    all_text = "\n".join(shape.text for shape in text_shapes) + "\n" + slide.notes_slide.notes_text_frame.text
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker.lower() in all_text.lower()]
    hygiene_passed = artifact_files == [OUTPUT_NAME] and not markers
    criteria.append(result_item("R010", hygiene_passed, f"输出文件={artifact_files}，可疑标记={markers}"))

    severe = not all(next(item for item in criteria if item["criterion_id"] == identifier)["raw"]["passed"] for identifier in ["R001", "R002", "R003", "R004", "R005", "R006", "R008", "R009"]) or bool(markers)
    criteria.append(result_item("R011", not severe, f"严重错误={severe}"))
    order = {identifier: index for index, identifier in enumerate(SPECS)}
    criteria.sort(key=lambda item: order[item["criterion_id"]])
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    output = args.artifact_dir / OUTPUT_NAME
    if output.is_file():
        evaluation = evaluate(output, args.artifact_dir)
    else:
        evaluation = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [], "criterion_results": [], "reward": 0.0, "passed": False, "hurdle_failures": ["missing_output"]}
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
