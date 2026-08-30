#!/usr/bin/env python3
"""PPTArena Case 47 专用 Golden 构建器：只转换冻结图表部件。"""

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from lxml import etree
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE


INPUT_SHA256 = "badd9134a75308729132505dc232413229bf20914fa0fddb42f4bf186b420114"
CHART_PART = "ppt/charts/chart1.xml"
C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
CATEGORIES = ("IALS 1994", "ALL 2003", "PIAAC 2012/2014")
VALUES = (273.475674150447, 268.246769360089, 271.722770340152)
CHART_GEOMETRY = (683418, 1836530, 7777163, 4133850)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def qn(namespace: str, tag: str) -> str:
    return f"{{{namespace}}}{tag}"


def chart_shape(presentation: Presentation):
    charts = [
        shape
        for shape in presentation.slides[0].shapes
        if shape.shape_type == MSO_SHAPE_TYPE.CHART
    ]
    if len(charts) != 1:
        raise ValueError(f"冻结文稿应有唯一图表，实际 {len(charts)}")
    return charts[0]


def audit_source(path: Path) -> None:
    if sha256(path) != INPUT_SHA256:
        raise ValueError("冻结输入 SHA256 不匹配")
    presentation = Presentation(path)
    if len(presentation.slides) != 1:
        raise ValueError("冻结输入不再是 1 页")
    shape = chart_shape(presentation)
    chart = shape.chart
    categories = tuple(str(value) for value in chart.plots[0].categories)
    values = tuple(float(value) for value in chart.series[0].values)
    if chart.chart_type != XL_CHART_TYPE.LINE_MARKERS:
        raise ValueError("冻结输入不再是带数据标记的折线图")
    if categories != CATEGORIES or values != VALUES or len(chart.series) != 1:
        raise ValueError("冻结图表的分类或数值发生变化")
    if tuple(int(value) for value in (shape.left, shape.top, shape.width, shape.height)) != CHART_GEOMETRY:
        raise ValueError("冻结图表几何发生变化")


def transform_chart_xml(data: bytes) -> bytes:
    root = etree.fromstring(data)
    line_chart = root.find(f".//{qn(C_NS, 'lineChart')}")
    if line_chart is None or root.find(f".//{qn(C_NS, 'barChart')}") is not None:
        raise ValueError("图表 XML 结构与本题冻结输入不符")
    line_chart.tag = qn(C_NS, "barChart")
    grouping = line_chart.find(qn(C_NS, "grouping"))
    if grouping is None:
        raise ValueError("图表缺少 grouping")
    grouping.set("val", "clustered")
    grouping.addprevious(etree.Element(qn(C_NS, "barDir"), val="col"))

    for series in line_chart.findall(qn(C_NS, "ser")):
        for tag in ("marker", "smooth"):
            element = series.find(qn(C_NS, tag))
            if element is not None:
                series.remove(element)
        shape_properties = series.find(qn(C_NS, "spPr"))
        if shape_properties is None:
            shape_properties = etree.Element(qn(C_NS, "spPr"))
            series.insert(2, shape_properties)
        for child in list(shape_properties):
            shape_properties.remove(child)
        solid = etree.SubElement(shape_properties, qn(A_NS, "solidFill"))
        etree.SubElement(solid, qn(A_NS, "schemeClr"), val="accent1")
        line = etree.SubElement(shape_properties, qn(A_NS, "ln"))
        etree.SubElement(line, qn(A_NS, "noFill"))
        shape_properties.addnext(etree.Element(qn(C_NS, "invertIfNegative"), val="0"))

    for tag in ("marker", "smooth"):
        element = line_chart.find(qn(C_NS, tag))
        if element is not None:
            line_chart.remove(element)
    first_axis = line_chart.find(qn(C_NS, "axId"))
    if first_axis is None:
        raise ValueError("图表缺少坐标轴关联")
    first_axis.addprevious(etree.Element(qn(C_NS, "gapWidth"), val="219"))
    first_axis.addprevious(etree.Element(qn(C_NS, "overlap"), val="-27"))
    return etree.tostring(
        root,
        xml_declaration=True,
        encoding="UTF-8",
        standalone=True,
    )


def build(source: Path, output: Path) -> None:
    audit_source(source)
    output.parent.mkdir(parents=True, exist_ok=True)
    changed = 0
    with ZipFile(source) as source_zip, ZipFile(output, "w", ZIP_DEFLATED) as output_zip:
        for item in source_zip.infolist():
            data = source_zip.read(item.filename)
            if item.filename == CHART_PART:
                data = transform_chart_xml(data)
                changed += 1
            output_zip.writestr(item, data)
    if changed != 1:
        raise ValueError("未唯一修改图表部件")
    validate_output(source, output)


def validate_output(source: Path, output: Path) -> None:
    presentation = Presentation(output)
    shape = chart_shape(presentation)
    chart = shape.chart
    if chart.chart_type != XL_CHART_TYPE.COLUMN_CLUSTERED:
        raise ValueError("产物不是簇状柱形图")
    if chart.has_legend or chart.has_title:
        raise ValueError("产物意外出现图例或图表标题")
    if tuple(str(value) for value in chart.plots[0].categories) != CATEGORIES:
        raise ValueError("产物分类不匹配")
    if tuple(float(value) for value in chart.series[0].values) != VALUES:
        raise ValueError("产物数值不匹配")
    if tuple(int(value) for value in (shape.left, shape.top, shape.width, shape.height)) != CHART_GEOMETRY:
        raise ValueError("产物图表几何不匹配")
    with ZipFile(source) as source_zip, ZipFile(output) as output_zip:
        source_names = source_zip.namelist()
        if source_names != output_zip.namelist():
            raise ValueError("产物包成员列表发生变化")
        changed = [
            name
            for name in source_names
            if source_zip.read(name) != output_zip.read(name)
        ]
    if changed != [CHART_PART]:
        raise ValueError(f"非预期包成员发生变化：{changed}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.source, args.output)
    print(f"{args.output}\nsha256={sha256(args.output)}")


if __name__ == "__main__":
    main()
