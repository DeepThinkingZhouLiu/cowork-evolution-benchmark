#!/usr/bin/env python3
"""只为 PPTArena Case 47 构造可回放的合法变体、负例和中等例。"""

from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches


OUTPUT_NAME = "ChangeLineGraphtoBar_GroundTruthA.pptx"


def output_path(root: Path, case: str) -> Path:
    target = root / case / OUTPUT_NAME
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def only_chart(presentation: Presentation):
    charts = [
        shape
        for shape in presentation.slides[0].shapes
        if shape.shape_type == MSO_SHAPE_TYPE.CHART
    ]
    if len(charts) != 1:
        raise ValueError(f"应有唯一图表，实际 {len(charts)}")
    return charts[0]


def valid_alternate(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    chart_shape = only_chart(presentation)
    chart_shape.left += Inches(0.08)
    chart_shape.width -= Inches(0.08)
    chart = chart_shape.chart
    gap_width = chart._element.xpath(".//c:barChart/c:gapWidth")
    if len(gap_width) != 1:
        raise ValueError("无法唯一定位柱间距")
    gap_width[0].set("val", "150")
    fills = chart._element.xpath(".//c:barChart/c:ser/c:spPr/a:solidFill/a:schemeClr")
    if len(fills) != 1:
        raise ValueError("无法唯一定位数据系列填充")
    fills[0].set("val", "accent2")
    presentation.save(output)


def mutate_chart_xml(golden: Path, output: Path, xpath: str, attribute: str | None = None, value: str = "") -> None:
    presentation = Presentation(golden)
    nodes = only_chart(presentation).chart._element.xpath(xpath)
    if len(nodes) != 1:
        raise ValueError(f"无法唯一定位：{xpath}，数量={len(nodes)}")
    if attribute is None:
        nodes[0].text = value
    else:
        nodes[0].set(attribute, value)
    presentation.save(output)


def no_labels(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    chart = only_chart(presentation).chart
    for node in chart._element.xpath(".//c:barChart//c:dLbls"):
        for label in list(node.xpath("./c:dLbl")):
            node.remove(label)
        for show_value in node.xpath("./c:showVal"):
            show_value.set("val", "0")
    presentation.save(output)


def changed_slide_number(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    matches = [
        shape
        for shape in presentation.slides[0].shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip() == "14"
    ]
    if len(matches) != 1:
        raise ValueError("无法唯一定位页码占位符")
    matches[0].text = "1"
    presentation.save(output)


def changed_title(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    title = presentation.slides[0].shapes.title
    title.text = title.text.replace("literacy score", "numeracy score")
    presentation.save(output)


def removed_logo(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    pictures = [shape for shape in presentation.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    if len(pictures) != 1:
        raise ValueError("无法唯一定位 AIR Logo")
    pictures[0]._element.getparent().remove(pictures[0]._element)
    presentation.save(output)


def moved_chart(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    only_chart(presentation).left += Inches(0.50)
    presentation.save(output)


def resized_chart(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    only_chart(presentation).width -= Inches(0.75)
    presentation.save(output)


def added_legend(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    only_chart(presentation).chart.has_legend = True
    presentation.save(output)


def removed_chart(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    chart = only_chart(presentation)
    chart._element.getparent().remove(chart._element)
    presentation.save(output)


def extra_chart(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    data = ChartData()
    data.categories = ["A", "B"]
    data.add_series("extra", [1, 2])
    presentation.slides[0].shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(7.8), Inches(5.6), Inches(1.4), Inches(0.9), data,
    )
    presentation.save(output)


def flattened(golden: Path, slide_image: Path, output: Path) -> None:
    presentation = Presentation(golden)
    slide = presentation.slides[0]
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    slide.shapes.add_picture(str(slide_image), 0, 0, width=presentation.slide_width, height=presentation.slide_height)
    presentation.save(output)


def hardcoded_marker(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    box = presentation.slides[0].shapes.add_textbox(Inches(0.10), Inches(7.10), Inches(2.50), Inches(0.20))
    box.text = "solution/Golden"
    presentation.save(output)


def axis_changed(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    only_chart(presentation).chart.value_axis.minimum_scale = 240.0
    presentation.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--slide-image", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()

    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))
    valid_alternate(args.golden, output_path(args.outdir, "valid-alternate"))
    mutate_chart_xml(args.golden, output_path(args.outdir, "horizontal-bar"), ".//c:barChart/c:barDir", "val", "bar")
    mutate_chart_xml(args.golden, output_path(args.outdir, "stacked-column"), ".//c:barChart/c:grouping", "val", "stacked")
    mutate_chart_xml(args.golden, output_path(args.outdir, "wrong-value"), ".//c:barChart/c:ser/c:val//c:pt[@idx='0']/c:v", None, "999")
    mutate_chart_xml(args.golden, output_path(args.outdir, "wrong-category"), ".//c:barChart/c:ser/c:cat//c:pt[@idx='2']/c:v", None, "Other")
    no_labels(args.golden, output_path(args.outdir, "no-labels"))
    added_legend(args.golden, output_path(args.outdir, "added-legend"))
    moved_chart(args.golden, output_path(args.outdir, "moved-chart"))
    resized_chart(args.golden, output_path(args.outdir, "resized-chart"))
    changed_slide_number(args.golden, output_path(args.outdir, "changed-slide-number"))
    changed_title(args.golden, output_path(args.outdir, "changed-title"))
    removed_logo(args.golden, output_path(args.outdir, "removed-logo"))
    removed_chart(args.golden, output_path(args.outdir, "removed-chart"))
    extra_chart(args.golden, output_path(args.outdir, "extra-chart"))
    flattened(args.golden, args.slide_image, output_path(args.outdir, "flattened"))
    hardcoded_marker(args.golden, output_path(args.outdir, "hardcoded-marker"))
    axis_changed(args.golden, output_path(args.outdir, "axis-changed"))


if __name__ == "__main__":
    main()
