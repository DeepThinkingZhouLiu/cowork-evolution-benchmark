#!/usr/bin/env python3
"""PPTArena Case 47 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import posixpath
import re
import statistics
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from PIL import Image, ImageChops, ImageDraw, ImageStat
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn


OUTPUT_NAME = "ChangeLineGraphtoBar_GroundTruthA.pptx"
INPUT_NAME = "ChangeLineGraphtoBar_TestA.pptx"
INPUT_SHA256 = "badd9134a75308729132505dc232413229bf20914fa0fddb42f4bf186b420114"
TASK_ID = "cowork-evo/pptarena-047-line-to-clustered-column-chart"
RUBRIC_VERSION = "pptarena-case-047-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-047-deterministic-v0.1"
EMU_PER_INCH = 914_400
CATEGORIES = ("IALS 1994", "ALL 2003", "PIAAC 2012/2014")
VALUES = (273.475674150447, 268.246769360089, 271.722770340152)
SUSPICIOUS_MARKERS = [
    "solution/golden", "judge-result", "artifact_dir", "pptarena-case-047", "/root/"
]

SPECS = {
    "R001": ("PPTX 完整且仍为原尺寸单页文稿", 8, "hurdle"),
    "R002": ("恰有一个原生可编辑的簇状柱形图", 18, "hurdle"),
    "R003": ("三个分类、顺序和精确数值保持", 20, "hurdle"),
    "R004": ("所有非图表对象和备注保持", 18, "hurdle"),
    "R005": ("无图例和图表标题，保留数据标签", 10, "hurdle"),
    "R006": ("图表几何近似原位置且页内无遮挡", 9, "hurdle"),
    "R007": ("坐标轴范围和图表可读性良好", 7, "positive"),
    "R008": ("可渲染，非图表区保持且图表非空白", 6, "hurdle"),
    "R009": ("无宏或损坏关系，输出目录干净", 4, "positive"),
    "R010": ("无类型、数据、保持或可编辑性严重错误", -20, "penalty"),
}


def normalize(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def locate_source() -> Path | None:
    candidates = [
        Path("/app/input_files") / INPUT_NAME,
        Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME,
    ]
    return next((path for path in candidates if path.is_file()), None)


def geometry(shape) -> tuple[int, int, int, int]:
    return int(shape.left), int(shape.top), int(shape.width), int(shape.height)


def rounded_geometry(shape) -> tuple[int, int, int, int]:
    return tuple(round(value / 1000) * 1000 for value in geometry(shape))


def safe(callback, default=""):
    try:
        return callback()
    except Exception:
        return default


def font_signature(font) -> tuple:
    color = font.color
    return (
        font.name,
        round(font.size.pt, 4) if font.size else None,
        font.bold,
        font.italic,
        font.underline,
        safe(lambda: str(font.language_id)),
        safe(lambda: str(color.type)),
        safe(lambda: str(color.rgb)),
        safe(lambda: str(color.theme_color)),
        safe(lambda: round(color.brightness, 5)),
    )


def text_signature(shape) -> tuple:
    if not getattr(shape, "has_text_frame", False):
        return ()
    frame = shape.text_frame
    paragraphs = tuple(
        (
            paragraph.text,
            paragraph.level,
            str(paragraph.alignment),
            str(paragraph.line_spacing),
            str(paragraph.space_before),
            str(paragraph.space_after),
            tuple((run.text, font_signature(run.font)) for run in paragraph.runs),
        )
        for paragraph in frame.paragraphs
    )
    return (
        normalize(shape.text),
        (
            frame.margin_left,
            frame.margin_right,
            frame.margin_top,
            frame.margin_bottom,
            str(frame.vertical_anchor),
            frame.word_wrap,
            str(frame.auto_size),
        ),
        paragraphs,
    )


def nonchart_signature(slide) -> tuple:
    signature = []
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.CHART:
            continue
        picture_hash = ""
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            picture_hash = hashlib.sha256(shape.image.blob).hexdigest()
        signature.append(
            (
                str(shape.shape_type),
                str(shape.placeholder_format.type) if shape.is_placeholder else "",
                rounded_geometry(shape),
                text_signature(shape),
                picture_hash,
            )
        )
    return tuple(signature)


def notes_signature(slide) -> tuple[str, ...]:
    try:
        return tuple(
            shape.text
            for shape in slide.notes_slide.shapes
            if getattr(shape, "has_text_frame", False)
        )
    except Exception:
        return ()


def charts(slide) -> list:
    return [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.CHART]


def chart_data(shape) -> tuple[tuple[str, ...], tuple[tuple[str, tuple[float | None, ...]], ...]]:
    chart = shape.chart
    categories = tuple(normalize(str(value)) for value in chart.plots[0].categories)
    series = tuple(
        (
            normalize(item.name or ""),
            tuple(None if value is None else float(value) for value in item.values),
        )
        for item in chart.series
    )
    return categories, series


def data_labels_visible(chart) -> tuple[bool, dict]:
    nodes = chart._element.xpath(".//c:barChart/c:dLbls | .//c:barChart/c:ser/c:dLbls")
    show_value = False
    custom_labels = 0
    for node in nodes:
        for element in node.xpath("./c:showVal"):
            show_value = show_value or str(element.get("val", "1")).lower() not in {"0", "false"}
        custom_labels += len(node.xpath("./c:dLbl[c:tx or c:txPr]"))
    passed = bool(nodes) and (show_value or custom_labels >= 3)
    return passed, {"dLbls_nodes": len(nodes), "show_value": show_value, "custom_labels": custom_labels}


def inside_slide(shape, slide_width: int, slide_height: int, tolerance_in: float = 0.03) -> bool:
    tolerance = int(tolerance_in * EMU_PER_INCH)
    left, top, width, height = geometry(shape)
    return (
        left >= -tolerance
        and top >= -tolerance
        and left + width <= slide_width + tolerance
        and top + height <= slide_height + tolerance
    )


def rectangles_overlap(first: tuple[int, int, int, int], second: tuple[int, int, int, int]) -> int:
    left1, top1, width1, height1 = first
    left2, top2, width2, height2 = second
    width = max(0, min(left1 + width1, left2 + width2) - max(left1, left2))
    height = max(0, min(top1 + height1, top2 + height2) - max(top1, top2))
    return width * height


def package_checks(path: Path) -> tuple[bool, dict]:
    evidence = {"broken_member": None, "missing_targets": [], "macros": []}
    try:
        with zipfile.ZipFile(path) as archive:
            evidence["broken_member"] = archive.testzip()
            names = set(archive.namelist())
            evidence["macros"] = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
            for rels_name in sorted(name for name in names if name.endswith(".rels")):
                root = __import__("xml.etree.ElementTree", fromlist=["ElementTree"]).fromstring(archive.read(rels_name))
                if rels_name == "_rels/.rels":
                    owner = ""
                else:
                    prefix, filename = rels_name.split("/_rels/", 1)
                    owner = posixpath.join(prefix, filename[:-5])
                base = posixpath.dirname(owner)
                for relation in root:
                    if relation.attrib.get("TargetMode") == "External":
                        continue
                    target = relation.attrib.get("Target", "")
                    resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                    if resolved not in names:
                        evidence["missing_targets"].append(f"{rels_name} -> {resolved}")
        passed = evidence["broken_member"] is None and not evidence["missing_targets"] and not evidence["macros"]
        return passed, evidence
    except Exception as exc:
        evidence["exception"] = str(exc)
        return False, evidence


def media_hashes(path: Path) -> Counter[str]:
    with zipfile.ZipFile(path) as archive:
        return Counter(
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("ppt/media/")
        )


def render(path: Path, outdir: Path) -> tuple[Path | None, int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "profile"
    profile.mkdir()
    completed = subprocess.run(
        ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(path)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return None, 0, completed.stdout[-500:]
    info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=30, check=False).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    pages = int(match.group(1)) if match else 0
    prefix = outdir / "slide"
    raster = subprocess.run(["pdftoppm", "-singlefile", "-png", "-r", "120", str(pdf), str(prefix)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=240, check=False)
    image = outdir / "slide.png"
    return image if raster.returncode == 0 and image.is_file() else None, pages, f"soffice={completed.returncode}, raster={raster.returncode}"


def render_metrics(source_image: Path, candidate_image: Path, chart_rect: tuple[int, int, int, int], slide_width: int, slide_height: int) -> tuple[float, float]:
    with Image.open(source_image) as source_raw, Image.open(candidate_image) as candidate_raw:
        source = source_raw.convert("RGB")
        candidate = candidate_raw.convert("RGB")
        if source.size != candidate.size:
            return 0.0, 0.0
        left, top, width, height = chart_rect
        box = (
            max(0, int(left * source.width / slide_width)),
            max(0, int(top * source.height / slide_height)),
            min(source.width, int(math.ceil((left + width) * source.width / slide_width))),
            min(source.height, int(math.ceil((top + height) * source.height / slide_height))),
        )
        mask = Image.new("L", source.size, 255)
        ImageDraw.Draw(mask).rectangle(box, fill=0)
        difference = ImageChops.difference(source, candidate)
        difference.paste((0, 0, 0), mask=ImageChops.invert(mask))
        nonchart_similarity = max(0.0, 1.0 - statistics.mean(ImageStat.Stat(difference).rms) / 255.0)
        chart_crop = candidate.crop(box).convert("L")
        histogram = chart_crop.histogram()
        dark_ratio = sum(histogram[:225]) / max(1, chart_crop.width * chart_crop.height)
        return nonchart_similarity, dark_ratio


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {
        "criterion_id": identifier,
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
        "raw": {"description": description, "weight": weight, "type": kind, "passed": bool(passed)},
    }


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [item["criterion_id"] for item in criteria if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0 and item["raw"]["passed"])
    penalty = sum(abs(item["raw"]["weight"]) for item in criteria if item["raw"]["weight"] < 0 and not item["raw"]["passed"])
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [sha256(path)] if path.is_file() else [],
        "criterion_results": criteria,
        "reward": reward,
        "passed": reward >= 0.8 and not hurdles,
        "hurdle_failures": hurdles,
    }


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        return make_result(path, [result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配") for identifier in SPECS])
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        return make_result(path, [result_item(identifier, False, f"PPTX 打开失败：{exc}") for identifier in SPECS])

    criteria: list[dict] = []
    package_passed, package_evidence = package_checks(path)
    size_passed = (
        len(source.slides) == 1
        and len(candidate.slides) == 1
        and source.slide_width == candidate.slide_width
        and source.slide_height == candidate.slide_height
    )
    criteria.append(result_item("R001", package_passed and size_passed, f"页数={len(candidate.slides)}，尺寸一致={size_passed}，包={json.dumps(package_evidence, ensure_ascii=False)}"))

    source_slide = source.slides[0]
    candidate_slide = candidate.slides[0]
    source_charts = charts(source_slide)
    candidate_charts = charts(candidate_slide)
    chart_type_passed = len(candidate_charts) == 1 and candidate_charts[0].chart.chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED
    full_slide_pictures = [shape.name for shape in candidate_slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.width * shape.height >= candidate.slide_width * candidate.slide_height * 0.80]
    criteria.append(result_item("R002", chart_type_passed and not full_slide_pictures, f"图表数={len(candidate_charts)}，类型={[str(shape.chart.chart_type) for shape in candidate_charts]}，整页图={full_slide_pictures}"))

    data_passed = False
    data_evidence = {}
    if len(candidate_charts) == 1:
        categories, series = chart_data(candidate_charts[0])
        values_match = (
            len(series) == 1
            and len(series[0][1]) == 3
            and all(actual is not None and math.isclose(actual, expected, rel_tol=1e-12, abs_tol=1e-12) for actual, expected in zip(series[0][1], VALUES, strict=True))
        )
        data_passed = categories == CATEGORIES and values_match and series[0][0] == ""
        data_evidence = {"categories": categories, "series": series, "values_match": values_match}
    criteria.append(result_item("R003", data_passed, json.dumps(data_evidence, ensure_ascii=False)))

    preservation_passed = (
        nonchart_signature(candidate_slide) == nonchart_signature(source_slide)
        and notes_signature(candidate_slide) == notes_signature(source_slide)
        and media_hashes(path) == media_hashes(source_path)
    )
    criteria.append(result_item("R004", preservation_passed, f"非图表签名一致={nonchart_signature(candidate_slide) == nonchart_signature(source_slide)}，备注一致={notes_signature(candidate_slide) == notes_signature(source_slide)}，媒体一致={media_hashes(path) == media_hashes(source_path)}"))

    labels_passed = False
    labels_evidence = {}
    if len(candidate_charts) == 1:
        chart = candidate_charts[0].chart
        visible, label_details = data_labels_visible(chart)
        labels_passed = not chart.has_legend and not chart.has_title and visible
        labels_evidence = {"legend": chart.has_legend, "title": chart.has_title, **label_details}
    criteria.append(result_item("R005", labels_passed, json.dumps(labels_evidence, ensure_ascii=False)))

    geometry_passed = False
    geometry_evidence = {}
    if len(source_charts) == 1 and len(candidate_charts) == 1:
        source_geometry = geometry(source_charts[0])
        candidate_geometry = geometry(candidate_charts[0])
        deltas = [abs(actual - expected) / EMU_PER_INCH for actual, expected in zip(candidate_geometry, source_geometry, strict=True)]
        overlaps = []
        for shape in candidate_slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.CHART:
                continue
            area = rectangles_overlap(candidate_geometry, geometry(shape)) / (EMU_PER_INCH**2)
            if area > 0.03:
                overlaps.append({"shape": shape.name, "area_in2": round(area, 4)})
        # 原图表与标题占位符边界有不到 0.01 in² 的四舍五入接触，不计为遮挡。
        geometry_passed = all(delta <= 0.25 for delta in deltas) and inside_slide(candidate_charts[0], candidate.slide_width, candidate.slide_height) and not overlaps
        geometry_evidence = {"source": source_geometry, "candidate": candidate_geometry, "deltas_in": [round(value, 4) for value in deltas], "overlaps": overlaps}
    criteria.append(result_item("R006", geometry_passed, json.dumps(geometry_evidence, ensure_ascii=False)))

    axis_passed = False
    axis_evidence = {}
    if len(candidate_charts) == 1:
        try:
            axis = candidate_charts[0].chart.value_axis
            axis_evidence = {"minimum": axis.minimum_scale, "maximum": axis.maximum_scale, "major": axis.major_unit}
            axis_passed = (
                math.isclose(float(axis.minimum_scale), 245.0, abs_tol=1e-9)
                and math.isclose(float(axis.maximum_scale), 285.0, abs_tol=1e-9)
                and math.isclose(float(axis.major_unit), 10.0, abs_tol=1e-9)
            )
        except Exception as exc:
            axis_evidence = {"exception": str(exc)}
    criteria.append(result_item("R007", axis_passed, json.dumps(axis_evidence, ensure_ascii=False)))

    prerequisites = all(next(item for item in criteria if item["criterion_id"] == identifier)["raw"]["passed"] for identifier in ("R001", "R002", "R003", "R004", "R005", "R006"))
    render_passed = False
    render_evidence = {"skipped": "前置门槛未通过"}
    if prerequisites:
        with tempfile.TemporaryDirectory(prefix="pptarena047-judge-") as temporary:
            root = Path(temporary)
            source_image, source_pages, source_note = render(source_path, root / "source")
            candidate_image, candidate_pages, candidate_note = render(path, root / "candidate")
            similarity, dark_ratio = (0.0, 0.0)
            if source_image and candidate_image:
                similarity, dark_ratio = render_metrics(source_image, candidate_image, geometry(candidate_charts[0]), candidate.slide_width, candidate.slide_height)
            render_passed = source_pages == 1 and candidate_pages == 1 and similarity >= 0.998 and dark_ratio >= 0.025
            render_evidence = {"source": source_note, "candidate": candidate_note, "source_pages": source_pages, "candidate_pages": candidate_pages, "nonchart_similarity": round(similarity, 6), "chart_dark_ratio": round(dark_ratio, 6)}
    criteria.append(result_item("R008", render_passed, json.dumps(render_evidence, ensure_ascii=False)))

    artifact_files = sorted(str(item.relative_to(artifact_dir)) for item in artifact_dir.rglob("*") if item.is_file())
    all_text = "\n".join(shape.text for shape in candidate_slide.shapes if getattr(shape, "has_text_frame", False)) + "\n" + "\n".join(notes_signature(candidate_slide))
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker.lower() in all_text.lower()]
    hygiene_passed = package_passed and artifact_files == [OUTPUT_NAME] and not markers
    criteria.append(result_item("R009", hygiene_passed, f"输出文件={artifact_files}，可疑标记={markers}，包完整={package_passed}"))

    severe = not all(next(item for item in criteria if item["criterion_id"] == identifier)["raw"]["passed"] for identifier in ("R001", "R002", "R003", "R004", "R005", "R006", "R008")) or bool(markers)
    criteria.append(result_item("R010", not severe, f"严重错误={severe}"))
    order = {identifier: index for index, identifier in enumerate(SPECS)}
    criteria.sort(key=lambda item: order[item["criterion_id"]])
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    output = args.artifact_dir / OUTPUT_NAME
    if output.is_file():
        evaluation = evaluate(output, args.artifact_dir)
    else:
        evaluation = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [], "criterion_results": [], "reward": 0.0, "passed": False, "hurdle_failures": ["missing_output"]}
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
