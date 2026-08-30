#!/usr/bin/env python3
"""PPTArena Case 50 专用 Golden 构建器：逐 run 统一锁定日期。"""

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

from pptx import Presentation


INPUT_SHA256 = "3f7fecef5d4f4c23b882124ef9d0fb79faa56d7fce53979e18ce5909f65c46f0"

# (slide number, shape id) -> (source text, normalized text, exact count)
EDITS = {
    (2, 7): [("25 February 2022", "February 25, 2022", 1)],
    (2, 9): [
        ("28 March 2022", "March 28, 2022", 1),
        ("02/26/2022", "February 26, 2022", 1),
    ],
    (3, 7): [("5 march 2022", "march 5, 2022", 1)],
    (3, 9): [
        ("03/28/2022", "March 28, 2022", 1),
        ("26 February 2022", "February 26, 2022", 1),
    ],
    (6, 9): [
        ("28 March 2022", "March 28, 2022", 1),
        ("8 March 2022", "March 8, 2022", 2),
    ],
    (7, 9): [
        ("24/02/2022", "February 24, 2022", 1),
        ("03/02/2022", "March 2, 2022", 1),
        ("effective as of  February 24.", "effective as of February 24, 2022.", 1),
        ("effective as of March 2.", "effective as of March 2, 2022.", 1),
        (
            "As of February 24 and March 2 the US",
            "As of February 24, 2022 and March 2, 2022, the US",
            1,
        ),
        ("As of March 2, no part", "As of March 2, 2022, no part", 1),
    ],
    (8, 3): [("22 March 2022", "March 22, 2022", 1)],
    (8, 9): [("Effective March 2, airlines", "Effective March 2, 2022 airlines", 1)],
    (9, 9): [("18 of March 2022", "March 18, 2022", 1)],
    (10, 7): [("12 March 2022", "March 12, 2022", 1)],
    (10, 9): [("March 12th 2022", "March 12, 2022", 2)],
    (11, 9): [("29 October 1998", "October 29, 1998", 1)],
    (12, 9): [
        ("1 September 2022", "September 1, 2022", 5),
        ("1 March 2022", "March 1, 2022", 5),
        ("first of September 2022", "September 1, 2022", 1),
        ("first of March 2022", "March 1, 2022", 1),
    ],
    (14, 9): [
        (
            "28 day of September – 10/12/1999",
            "September 28, 1999 - October 12, 1999",
            1,
        )
    ],
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def shape_by_id(presentation: Presentation, slide_number: int, shape_id: int):
    matches = [
        shape
        for shape in presentation.slides[slide_number - 1].shapes
        if shape.shape_id == shape_id
    ]
    if len(matches) != 1 or not getattr(matches[0], "has_text_frame", False):
        raise ValueError(f"无法唯一定位文本对象：slide={slide_number}, shape={shape_id}")
    return matches[0]


def plain_transform(text: str, replacements: list[tuple[str, str, int]]) -> str:
    for old, new, expected_count in replacements:
        actual = text.count(old)
        if actual != expected_count:
            raise ValueError(f"来源日期计数错误：{old!r}，期望 {expected_count}，实际 {actual}")
        text = text.replace(old, new)
    return text


def replace_in_paragraph(paragraph, old: str, new: str) -> int:
    runs = list(paragraph.runs)
    text = "".join(run.text for run in runs)
    positions = []
    start = 0
    while True:
        index = text.find(old, start)
        if index < 0:
            break
        positions.append(index)
        start = index + len(old)
    for index in reversed(positions):
        end = index + len(old)
        spans = []
        cursor = 0
        for run_index, run in enumerate(runs):
            next_cursor = cursor + len(run.text)
            spans.append((run_index, cursor, next_cursor))
            cursor = next_cursor
        start_info = next((item for item in spans if item[1] <= index < item[2]), None)
        end_info = next((item for item in spans if item[1] < end <= item[2]), None)
        if start_info is None or end_info is None:
            raise ValueError(f"无法将日期映射到 run：{old!r}")
        start_run = runs[start_info[0]]
        end_run = runs[end_info[0]]
        start_offset = index - start_info[1]
        end_offset = end - end_info[1]
        if start_info[0] == end_info[0]:
            start_run.text = start_run.text[:start_offset] + new + start_run.text[end_offset:]
        else:
            start_run.text = start_run.text[:start_offset] + new
            for run_index in range(start_info[0] + 1, end_info[0]):
                runs[run_index].text = ""
            end_run.text = end_run.text[end_offset:]
        text = "".join(run.text for run in runs)
    return len(positions)


def replace_in_shape(shape, old: str, new: str, expected_count: int) -> None:
    actual = sum(replace_in_paragraph(paragraph, old, new) for paragraph in shape.text_frame.paragraphs)
    if actual != expected_count:
        raise ValueError(f"run 编辑计数错误：{old!r}，期望 {expected_count}，实际 {actual}")


def shape_structure(presentation: Presentation) -> tuple:
    return tuple(
        tuple(
            (
                int(shape.shape_id),
                str(shape.shape_type),
                int(shape.left),
                int(shape.top),
                int(shape.width),
                int(shape.height),
                str(shape.placeholder_format.type) if shape.is_placeholder else "",
            )
            for shape in slide.shapes
        )
        for slide in presentation.slides
    )


def build(source: Path, output: Path) -> None:
    if sha256(source) != INPUT_SHA256:
        raise ValueError("冻结输入 SHA256 不匹配")
    source_presentation = Presentation(source)
    if len(source_presentation.slides) != 15:
        raise ValueError("冻结输入不再是 15 页")
    presentation = Presentation(source)
    expected_texts = {}
    for key, replacements in EDITS.items():
        shape = shape_by_id(presentation, *key)
        source_text = shape.text
        expected_texts[key] = plain_transform(source_text, replacements)
        for old, new, expected_count in replacements:
            replace_in_shape(shape, old, new, expected_count)
        if shape.text != expected_texts[key]:
            raise ValueError(f"日期编辑后文本不匹配：{key}")

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    validate_output(source_presentation, output, expected_texts)


def validate_output(source: Presentation, output: Path, expected_texts: dict) -> None:
    candidate = Presentation(output)
    if len(candidate.slides) != 15 or (candidate.slide_width, candidate.slide_height) != (source.slide_width, source.slide_height):
        raise ValueError("产物页数或页面尺寸不匹配")
    if shape_structure(candidate) != shape_structure(source):
        raise ValueError("产物对象数量、类型、占位符或几何发生变化")
    for slide_number, (source_slide, candidate_slide) in enumerate(zip(source.slides, candidate.slides, strict=True), 1):
        source_map = {shape.shape_id: shape for shape in source_slide.shapes}
        candidate_map = {shape.shape_id: shape for shape in candidate_slide.shapes}
        for shape_id, source_shape in source_map.items():
            if not getattr(source_shape, "has_text_frame", False):
                continue
            expected = expected_texts.get((slide_number, shape_id), source_shape.text)
            if candidate_map[shape_id].text != expected:
                raise ValueError(f"产物文本超出锁定日期编辑范围：{slide_number}/{shape_id}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.source, args.output)
    print(f"{args.output}\nsha256={sha256(args.output)}")


if __name__ == "__main__":
    main()
