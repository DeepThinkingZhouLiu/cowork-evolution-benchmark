#!/usr/bin/env python3
"""PPTArena Case 50 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import posixpath
import re
import statistics
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from PIL import Image, ImageChops, ImageDraw, ImageStat
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "NormalizeDate_GroundTruth.pptx"
INPUT_NAME = "NormalizeDate_Test.pptx"
INPUT_SHA256 = "3f7fecef5d4f4c23b882124ef9d0fb79faa56d7fce53979e18ce5909f65c46f0"
TASK_ID = "cowork-evo/pptarena-050-normalize-date-formats"
RUBRIC_VERSION = "pptarena-case-050-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-050-deterministic-v0.1"
DATE_INSTANCE_COUNT = 37
SUSPICIOUS_MARKERS = ["solution/golden", "judge-result", "artifact_dir", "pptarena-case-050", "/root/"]

EDITS = {
    (2, 7): [("25 February 2022", "February 25, 2022", 1)],
    (2, 9): [("28 March 2022", "March 28, 2022", 1), ("02/26/2022", "February 26, 2022", 1)],
    (3, 7): [("5 march 2022", "march 5, 2022", 1)],
    (3, 9): [("03/28/2022", "March 28, 2022", 1), ("26 February 2022", "February 26, 2022", 1)],
    (6, 9): [("28 March 2022", "March 28, 2022", 1), ("8 March 2022", "March 8, 2022", 2)],
    (7, 9): [
        ("24/02/2022", "February 24, 2022", 1),
        ("03/02/2022", "March 2, 2022", 1),
        ("effective as of  February 24.", "effective as of February 24, 2022.", 1),
        ("effective as of March 2.", "effective as of March 2, 2022.", 1),
        ("As of February 24 and March 2 the US", "As of February 24, 2022 and March 2, 2022, the US", 1),
        ("As of March 2, no part", "As of March 2, 2022, no part", 1),
    ],
    (8, 3): [("22 March 2022", "March 22, 2022", 1)],
    (8, 9): [("Effective March 2, airlines", "Effective March 2, 2022 airlines", 1)],
    (9, 9): [("18 of March 2022", "March 18, 2022", 1)],
    (10, 7): [("12 March 2022", "March 12, 2022", 1)],
    (10, 9): [("March 12th 2022", "March 12, 2022", 2)],
    (11, 9): [("29 October 1998", "October 29, 1998", 1)],
    (12, 9): [
        ("1 September 2022", "September 1, 2022", 5),
        ("1 March 2022", "March 1, 2022", 5),
        ("first of September 2022", "September 1, 2022", 1),
        ("first of March 2022", "March 1, 2022", 1),
    ],
    (14, 9): [("28 day of September – 10/12/1999", "September 28, 1999 - October 12, 1999", 1)],
}

SPECS = {
    "R001": ("PPTX 完整且仍为原尺寸 15 页文稿", 7, "hurdle"),
    "R002": ("37 个目标日期均已正确统一", 23, "hurdle"),
    "R003": ("缺失年份的日期与 1999 范围推断正确", 12, "hurdle"),
    "R004": ("原月份大小写保持", 7, "hurdle"),
    "R005": ("非日期文字、顺序、结尾页、页码页脚、备注图片保持", 16, "hurdle"),
    "R006": ("对象、几何、段落与逐字符样式保持", 14, "hurdle"),
    "R007": ("文本原生可编辑且无增删压平", 8, "hurdle"),
    "R008": ("15 页可渲染，日期区可见且其他区域保持", 8, "hurdle"),
    "R009": ("无宏、损坏关系、媒体丢失或答案标记", 5, "positive"),
    "R010": ("无日期、年份、保持、样式或可编辑性严重错误", -20, "penalty"),
}


def normalize(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def locate_source() -> Path | None:
    candidates = [Path("/app/input_files") / INPUT_NAME, Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME]
    return next((path for path in candidates if path.is_file()), None)


def safe(callback, default=""):
    try:
        return callback()
    except Exception:
        return default


def font_signature(font) -> tuple:
    color = font.color
    return (
        font.name,
        round(font.size.pt, 4) if font.size else None,
        font.bold,
        font.italic,
        font.underline,
        safe(lambda: str(font.language_id)),
        safe(lambda: str(color.type)),
        safe(lambda: str(color.rgb)),
        safe(lambda: str(color.theme_color)),
        safe(lambda: round(color.brightness, 5)),
    )


def paragraph_signature(paragraph) -> tuple:
    return (
        paragraph.level,
        str(paragraph.alignment),
        str(paragraph.line_spacing),
        str(paragraph.space_before),
        str(paragraph.space_after),
    )


def character_styles(paragraph) -> tuple:
    return tuple(style for run in paragraph.runs for style in [font_signature(run.font)] * len(run.text))


def text_frame_signature(shape) -> tuple:
    frame = shape.text_frame
    return (
        frame.margin_left,
        frame.margin_right,
        frame.margin_top,
        frame.margin_bottom,
        str(frame.vertical_anchor),
        frame.word_wrap,
        str(frame.auto_size),
    )


def shape_structure(presentation: Presentation) -> tuple:
    return tuple(
        tuple(
            (
                int(shape.shape_id), str(shape.shape_type), int(shape.left), int(shape.top), int(shape.width), int(shape.height),
                str(shape.placeholder_format.type) if shape.is_placeholder else "",
            )
            for shape in slide.shapes
        )
        for slide in presentation.slides
    )


def shape_map(presentation: Presentation, slide_number: int) -> dict:
    return {shape.shape_id: shape for shape in presentation.slides[slide_number - 1].shapes}


def plain_transform(text: str, replacements: list[tuple[str, str, int]]) -> str:
    for old, new, expected_count in replacements:
        if text.count(old) != expected_count:
            raise ValueError(f"冻结来源计数不匹配：{old!r}")
        text = text.replace(old, new)
    return text


def transform_chars(text: str, styles: tuple, replacements: list[tuple[str, str, int]]) -> tuple[str, tuple]:
    style_list = list(styles)
    for old, new, expected_count in replacements:
        positions = []
        start = 0
        while True:
            index = text.find(old, start)
            if index < 0:
                break
            positions.append(index)
            start = index + len(old)
        if len(positions) != expected_count:
            raise ValueError(f"冻结段落计数不匹配：{old!r}")
        for index in reversed(positions):
            end = index + len(old)
            replacement_style = style_list[index]
            text = text[:index] + new + text[end:]
            style_list[index:end] = [replacement_style] * len(new)
    return text, tuple(style_list)


def expected_paragraphs(source_shape, replacements: list[tuple[str, str, int]]) -> tuple:
    paragraphs = []
    remaining = {old: count for old, _, count in replacements}
    for paragraph in source_shape.text_frame.paragraphs:
        text = paragraph.text
        styles = character_styles(paragraph)
        for old, new, _ in replacements:
            count = text.count(old)
            if count:
                remaining[old] -= count
                text, styles = transform_chars(text, styles, [(old, new, count)])
        paragraphs.append((text, styles, paragraph_signature(paragraph)))
    if any(value != 0 for value in remaining.values()):
        raise ValueError(f"日期跨段落或冻结计数异常：{remaining}")
    return tuple(paragraphs)


def actual_paragraphs(shape) -> tuple:
    return tuple((paragraph.text, character_styles(paragraph), paragraph_signature(paragraph)) for paragraph in shape.text_frame.paragraphs)


def notes_signature(slide) -> tuple[str, ...]:
    try:
        return tuple(shape.text for shape in slide.notes_slide.shapes if getattr(shape, "has_text_frame", False))
    except Exception:
        return ()


def media_hashes(path: Path) -> Counter[str]:
    with zipfile.ZipFile(path) as archive:
        return Counter(hashlib.sha256(archive.read(name)).hexdigest() for name in archive.namelist() if name.startswith("ppt/media/"))


def package_checks(path: Path) -> tuple[bool, dict]:
    evidence = {"broken_member": None, "missing_targets": [], "macros": []}
    try:
        with zipfile.ZipFile(path) as archive:
            evidence["broken_member"] = archive.testzip()
            names = set(archive.namelist())
            evidence["macros"] = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
            for rels_name in sorted(name for name in names if name.endswith(".rels")):
                root = __import__("xml.etree.ElementTree", fromlist=["ElementTree"]).fromstring(archive.read(rels_name))
                if rels_name == "_rels/.rels":
                    owner = ""
                else:
                    prefix, filename = rels_name.split("/_rels/", 1)
                    owner = posixpath.join(prefix, filename[:-5])
                base = posixpath.dirname(owner)
                for relation in root:
                    if relation.attrib.get("TargetMode") == "External":
                        continue
                    target = relation.attrib.get("Target", "")
                    resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                    if resolved not in names:
                        evidence["missing_targets"].append(f"{rels_name} -> {resolved}")
        return evidence["broken_member"] is None and not evidence["missing_targets"] and not evidence["macros"], evidence
    except Exception as exc:
        evidence["exception"] = str(exc)
        return False, evidence


def render(path: Path, outdir: Path) -> tuple[list[Path], int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "profile"
    profile.mkdir()
    completed = subprocess.run(
        ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(path)],
        stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300, check=False,
    )
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return [], 0, completed.stdout[-500:]
    info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=30, check=False).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    pages = int(match.group(1)) if match else 0
    prefix = outdir / "slide"
    raster = subprocess.run(["pdftoppm", "-png", "-r", "80", str(pdf), str(prefix)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300, check=False)
    images = sorted(outdir.glob("slide-*.png")) if raster.returncode == 0 else []
    return images, pages, f"soffice={completed.returncode}, raster={raster.returncode}"


def render_metrics(source_images: list[Path], candidate_images: list[Path], source: Presentation, candidate: Presentation) -> tuple[float, list[float]]:
    similarities = []
    dark_ratios = []
    target_by_slide: dict[int, list] = {}
    for slide_number, shape_id in EDITS:
        target_by_slide.setdefault(slide_number, []).append(shape_map(candidate, slide_number)[shape_id])
    for slide_number, (source_path, candidate_path) in enumerate(zip(source_images, candidate_images, strict=True), 1):
        with Image.open(source_path) as source_raw, Image.open(candidate_path) as candidate_raw:
            first = source_raw.convert("RGB")
            second = candidate_raw.convert("RGB")
            if first.size != second.size:
                similarities.append(0.0)
                continue
            mask = Image.new("L", first.size, 255)
            draw = ImageDraw.Draw(mask)
            for shape in target_by_slide.get(slide_number, []):
                box = (
                    max(0, int(shape.left * first.width / candidate.slide_width)),
                    max(0, int(shape.top * first.height / candidate.slide_height)),
                    min(first.width, int(math.ceil((shape.left + shape.width) * first.width / candidate.slide_width))),
                    min(first.height, int(math.ceil((shape.top + shape.height) * first.height / candidate.slide_height))),
                )
                draw.rectangle(box, fill=0)
                crop = second.crop(box).convert("L")
                histogram = crop.histogram()
                dark_ratios.append(sum(histogram[:210]) / max(1, crop.width * crop.height))
            difference = ImageChops.difference(first, second)
            difference.paste((0, 0, 0), mask=ImageChops.invert(mask))
            similarities.append(max(0.0, 1.0 - statistics.mean(ImageStat.Stat(difference).rms) / 255.0))
    return min(similarities, default=0.0), dark_ratios


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {"criterion_id": identifier, "score": 1.0 if passed else 0.0, "evidence": evidence, "method": "deterministic", "raw": {"description": description, "weight": weight, "type": kind, "passed": bool(passed)}}


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [item["criterion_id"] for item in criteria if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0 and item["raw"]["passed"])
    penalty = sum(abs(item["raw"]["weight"]) for item in criteria if item["raw"]["weight"] < 0 and not item["raw"]["passed"])
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [sha256(path)] if path.is_file() else [], "criterion_results": criteria, "reward": reward, "passed": reward >= 0.8 and not hurdles, "hurdle_failures": hurdles}


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        return make_result(path, [result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配") for identifier in SPECS])
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        return make_result(path, [result_item(identifier, False, f"PPTX 打开失败：{exc}") for identifier in SPECS])

    criteria = []
    package_passed, package_evidence = package_checks(path)
    size_passed = len(source.slides) == len(candidate.slides) == 15 and (source.slide_width, source.slide_height) == (candidate.slide_width, candidate.slide_height)
    criteria.append(result_item("R001", package_passed and size_passed, f"页数={len(candidate.slides)}，尺寸一致={size_passed}，包={json.dumps(package_evidence, ensure_ascii=False)}"))

    expected_texts = {}
    target_results = []
    styles_passed = True
    target_text_passed = True
    for key, replacements in EDITS.items():
        source_shape = shape_map(source, key[0]).get(key[1])
        candidate_shape = shape_map(candidate, key[0]).get(key[1])
        if source_shape is None or candidate_shape is None or not getattr(candidate_shape, "has_text_frame", False):
            target_text_passed = styles_passed = False
            target_results.append({"target": key, "present": False})
            continue
        expected = plain_transform(source_shape.text, replacements)
        expected_texts[key] = expected
        text_ok = candidate_shape.text == expected
        try:
            style_ok = actual_paragraphs(candidate_shape) == expected_paragraphs(source_shape, replacements) and text_frame_signature(candidate_shape) == text_frame_signature(source_shape)
        except Exception:
            style_ok = False
        target_text_passed = target_text_passed and text_ok
        styles_passed = styles_passed and style_ok
        target_results.append({"target": key, "text_ok": text_ok, "style_ok": style_ok})
    criteria.append(result_item("R002", target_text_passed, f"日期实例={DATE_INSTANCE_COUNT}，目标={json.dumps(target_results, ensure_ascii=False)}"))

    combined_text = "\n".join(shape.text for slide in candidate.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False))
    inference_needles = [
        "February 24, 2022", "March 2, 2022", "Effective March 2, 2022 airlines",
        "September 28, 1999 - October 12, 1999",
    ]
    inference_passed = target_text_passed and all(needle in combined_text for needle in inference_needles)
    criteria.append(result_item("R003", inference_passed, f"上下文日期={[(needle, combined_text.count(needle)) for needle in inference_needles]}"))

    casing_passed = "march 5, 2022" in combined_text and "MARCH 11, 2022" in combined_text and "March 5, 2022" not in combined_text
    criteria.append(result_item("R004", casing_passed, "小写 march 5, 2022=%s，大写 MARCH 11, 2022=%s" % (combined_text.count("march 5, 2022"), combined_text.count("MARCH 11, 2022"))))

    structure_equal = shape_structure(candidate) == shape_structure(source)
    non_target_text_ok = True
    notes_ok = True
    for slide_number, (source_slide, candidate_slide) in enumerate(zip(source.slides, candidate.slides, strict=False), 1):
        source_shapes = {shape.shape_id: shape for shape in source_slide.shapes}
        candidate_shapes = {shape.shape_id: shape for shape in candidate_slide.shapes}
        for shape_id, source_shape in source_shapes.items():
            if (slide_number, shape_id) in EDITS or not getattr(source_shape, "has_text_frame", False):
                continue
            candidate_shape = candidate_shapes.get(shape_id)
            non_target_text_ok = non_target_text_ok and candidate_shape is not None and getattr(candidate_shape, "has_text_frame", False) and candidate_shape.text == source_shape.text
        notes_ok = notes_ok and notes_signature(candidate_slide) == notes_signature(source_slide)
    media_ok = media_hashes(path) == media_hashes(source_path)
    criteria.append(result_item("R005", structure_equal and non_target_text_ok and notes_ok and media_ok, f"结构={structure_equal}，非日期文本={non_target_text_ok}，备注={notes_ok}，媒体={media_ok}，第15页对象={len(candidate.slides[14].shapes) if len(candidate.slides) >= 15 else -1}"))

    all_styles_ok = styles_passed
    if structure_equal:
        for slide_number, (source_slide, candidate_slide) in enumerate(zip(source.slides, candidate.slides, strict=True), 1):
            candidate_shapes = {shape.shape_id: shape for shape in candidate_slide.shapes}
            for source_shape in source_slide.shapes:
                if not getattr(source_shape, "has_text_frame", False) or (slide_number, source_shape.shape_id) in EDITS:
                    continue
                candidate_shape = candidate_shapes[source_shape.shape_id]
                all_styles_ok = all_styles_ok and actual_paragraphs(candidate_shape) == actual_paragraphs(source_shape) and text_frame_signature(candidate_shape) == text_frame_signature(source_shape)
    else:
        all_styles_ok = False
    criteria.append(result_item("R006", all_styles_ok, f"目标样式={styles_passed}，全部样式={all_styles_ok}"))

    page_area = candidate.slide_width * candidate.slide_height
    full_slide_pictures = [f"{i}:{shape.name}" for i, slide in enumerate(candidate.slides, 1) for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.width * shape.height >= page_area * 0.80]
    editable_passed = structure_equal and target_text_passed and not full_slide_pictures and all(shape_map(candidate, key[0]).get(key[1]) is not None for key in EDITS)
    criteria.append(result_item("R007", editable_passed, f"结构={structure_equal}，目标文本对象=14/14，整页图={full_slide_pictures}"))

    prerequisites = all(next(item for item in criteria if item["criterion_id"] == identifier)["raw"]["passed"] for identifier in ("R001", "R002", "R003", "R004", "R005", "R006", "R007"))
    render_passed = False
    render_evidence = {"skipped": "前置门槛未通过"}
    if prerequisites:
        with tempfile.TemporaryDirectory(prefix="pptarena050-judge-") as temporary:
            root = Path(temporary)
            source_images, source_pages, source_note = render(source_path, root / "source")
            candidate_images, candidate_pages, candidate_note = render(path, root / "candidate")
            similarity, dark_ratios = (0.0, [])
            if len(source_images) == len(candidate_images) == 15:
                similarity, dark_ratios = render_metrics(source_images, candidate_images, source, candidate)
            render_passed = source_pages == candidate_pages == 15 and len(dark_ratios) == 14 and min(dark_ratios, default=0.0) >= 0.0015 and similarity >= 0.999
            render_evidence = {"source": source_note, "candidate": candidate_note, "source_pages": source_pages, "candidate_pages": candidate_pages, "non_target_min_similarity": round(similarity, 6), "target_dark_ratio_min": round(min(dark_ratios, default=0.0), 6), "target_regions": len(dark_ratios)}
    criteria.append(result_item("R008", render_passed, json.dumps(render_evidence, ensure_ascii=False)))

    artifact_files = sorted(str(item.relative_to(artifact_dir)) for item in artifact_dir.rglob("*") if item.is_file())
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker.lower() in combined_text.lower()]
    hygiene_passed = package_passed and media_ok and artifact_files == [OUTPUT_NAME] and not markers
    criteria.append(result_item("R009", hygiene_passed, f"输出={artifact_files}，可疑标记={markers}，媒体={media_ok}"))

    severe = not all(next(item for item in criteria if item["criterion_id"] == identifier)["raw"]["passed"] for identifier in ("R001", "R002", "R003", "R004", "R005", "R006", "R007", "R008")) or bool(markers)
    criteria.append(result_item("R010", not severe, f"严重错误={severe}"))
    order = {identifier: index for index, identifier in enumerate(SPECS)}
    criteria.sort(key=lambda item: order[item["criterion_id"]])
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    output = args.artifact_dir / OUTPUT_NAME
    if output.is_file():
        evaluation = evaluate(output, args.artifact_dir)
    else:
        evaluation = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [], "criterion_results": [], "reward": 0.0, "passed": False, "hurdle_failures": ["missing_output"]}
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
