#!/usr/bin/env python3
"""从 PPTArena Case 53 冻结输入构建经逐对象审计的本地 Golden。"""

from __future__ import annotations

import argparse
import hashlib
import zipfile
from collections import Counter
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt


INPUT_SHA256 = "69eb80458f9c78ee867e27812af34f63d12483fb1e4b4a9dd3eb63c3cdc37f99"
SOURCE_TEXT = (
    "Revenue 30% AMI Rents Annual Actualized Income $387,030 "
    "Annual Operating Expenses ($420,300) Annual Net Revenue ($33,270)"
)
VALUES = (
    ("Revenue", "30% AMI Rents"),
    ("Annual Actualized Income", "$387,030"),
    ("Annual Operating Expenses", "($420,300)"),
    ("Annual Net Revenue", "($33,270)"),
)
FONT_SIZES = ((28, 16), (16, 16), (16, 16), (18, 18))
BOLD = ((True, False), (False, False), (True, False), (False, True))
TARGET_TOP = 3_634_201


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def media_hashes(path: Path) -> Counter[str]:
    with zipfile.ZipFile(path) as archive:
        return Counter(
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("ppt/media/")
        )


def shape_by_id(presentation: Presentation, shape_id: int):
    matches = [shape for shape in presentation.slides[0].shapes if shape.shape_id == shape_id]
    if len(matches) != 1:
        raise ValueError(f"无法唯一定位对象 {shape_id}")
    return matches[0]


def table_style_state(shape) -> tuple:
    return (
        tuple(row.height for row in shape.table.rows),
        tuple(column.width for column in shape.table.columns),
        etree.tostring(shape._element.graphic.graphicData.tbl.tblPr),
        tuple(
            etree.tostring(cell._tc.get_or_add_tcPr())
            for row in shape.table.rows
            for cell in row.cells
        ),
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    if not args.source.is_file() or sha256(args.source) != INPUT_SHA256:
        raise ValueError("冻结输入缺失或 SHA256 不匹配")

    presentation = Presentation(args.source)
    if len(presentation.slides) != 1 or len(presentation.slides[0].shapes) != 9:
        raise ValueError("冻结页面结构已变化")

    source_box = shape_by_id(presentation, 9)
    target = shape_by_id(presentation, 16)
    if not source_box.has_text_frame or source_box.text != SOURCE_TEXT:
        raise ValueError("源数据文本框内容已变化")
    if not target.has_table or len(target.table.rows) != 4 or len(target.table.columns) != 2:
        raise ValueError("目标表格结构已变化")
    if any(cell.text for row in target.table.rows for cell in row.cells):
        raise ValueError("目标表格不再为空")

    style_before = table_style_state(target)
    media_before = media_hashes(args.source)
    other_shapes_before = {
        shape.shape_id: (str(shape.shape_type), shape.left, shape.top, shape.width, shape.height, getattr(shape, "text", ""))
        for shape in presentation.slides[0].shapes
        if shape.shape_id not in {9, 16}
    }

    source_box._element.getparent().remove(source_box._element)
    target.top = TARGET_TOP
    for row_index, row in enumerate(target.table.rows):
        for column_index, cell in enumerate(row.cells):
            cell.text = VALUES[row_index][column_index]
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.LEFT if column_index == 0 else PP_ALIGN.RIGHT
            run = paragraph.runs[0]
            run.font.name = "Overpass Medium"
            run.font.size = Pt(FONT_SIZES[row_index][column_index])
            run.font.bold = BOLD[row_index][column_index]
            run.font.italic = False
            if (row_index, column_index) == (3, 1):
                run.font.color.theme_color = MSO_THEME_COLOR.ACCENT_5
            else:
                run.font.color.rgb = RGBColor(0, 0, 0)

    if table_style_state(target) != style_before:
        raise ValueError("填入文本时破坏了表格边框、填充、行高、列宽或主题样式")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(args.output)

    rebuilt = Presentation(args.output)
    rebuilt_target = shape_by_id(rebuilt, 16)
    actual = tuple(tuple(cell.text for cell in row.cells) for row in rebuilt_target.table.rows)
    if actual != VALUES:
        raise ValueError(f"Golden 表格真值回读失败：{actual!r}")
    if any(shape.shape_id == 9 for shape in rebuilt.slides[0].shapes):
        raise ValueError("源数据文本框未删除")
    if len(rebuilt.slides[0].shapes) != 8 or len([shape for shape in rebuilt.slides[0].shapes if shape.has_table]) != 3:
        raise ValueError("输出增删了任务边界之外的对象或表格")
    if table_style_state(rebuilt_target) != style_before:
        raise ValueError("保存后表格底层样式已变化")
    other_shapes_after = {
        shape.shape_id: (str(shape.shape_type), shape.left, shape.top, shape.width, shape.height, getattr(shape, "text", ""))
        for shape in rebuilt.slides[0].shapes
        if shape.shape_id != 16
    }
    if other_shapes_after != other_shapes_before:
        raise ValueError("非目标对象文本或几何已变化")
    if media_hashes(args.output) != media_before:
        raise ValueError("媒体资产已变化")

    print(f"{args.output} {sha256(args.output)}")


if __name__ == "__main__":
    main()
