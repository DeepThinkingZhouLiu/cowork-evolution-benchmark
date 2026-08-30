#!/usr/bin/env python3
"""只为 PPTArena Case 53 构造合法变体与数据/表格/保持负例。"""

from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUTPUT_NAME = "FillTablefromData_GroundTruthA.pptx"
SOURCE_TEXT = (
    "Revenue 30% AMI Rents Annual Actualized Income $387,030 "
    "Annual Operating Expenses ($420,300) Annual Net Revenue ($33,270)"
)


def output_path(root: Path, case: str) -> Path:
    target = root / case / OUTPUT_NAME
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def shape_by_id(presentation: Presentation, shape_id: int):
    matches = [shape for shape in presentation.slides[0].shapes if shape.shape_id == shape_id]
    if len(matches) != 1:
        raise ValueError(f"无法唯一定位对象 {shape_id}")
    return matches[0]


def target_table(presentation: Presentation):
    matches = [shape for shape in presentation.slides[0].shapes if shape.has_table and len(shape.table.rows) == 4 and len(shape.table.columns) == 2]
    if len(matches) != 1:
        raise ValueError("无法唯一定位 4x2 目标表格")
    return matches[0]


def valid_alternate(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    target = target_table(presentation)
    target.top = Inches(4.25)
    for row in target.table.rows:
        cell = row.cells[1]
        for paragraph in cell.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(17 if cell.text != "($33,270)" else 18)
    presentation.save(output)


def wrong_value(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    target_table(presentation).table.cell(1, 1).text = "$386,030"
    presentation.save(output)


def swapped_columns(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    table = target_table(presentation).table
    left, right = table.cell(0, 0).text, table.cell(0, 1).text
    table.cell(0, 0).text, table.cell(0, 1).text = right, left
    presentation.save(output)


def kept_source_textbox(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    slide = presentation.slides[0]
    box = slide.shapes.add_textbox(7_433_010, 3_429_000, 4_128_091, 1_200_329)
    box.text = SOURCE_TEXT
    presentation.save(output)


def fake_textbox_table(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    slide = presentation.slides[0]
    target = target_table(presentation)
    values = [[cell.text for cell in row.cells] for row in target.table.rows]
    left, top = target.left, target.top
    column_widths = [column.width for column in target.table.columns]
    row_heights = [row.height for row in target.table.rows]
    target._element.getparent().remove(target._element)
    y = top
    for row_index, height in enumerate(row_heights):
        x = left
        for column_index, width in enumerate(column_widths):
            box = slide.shapes.add_textbox(x, y, width, height)
            box.text = values[row_index][column_index]
            x += width
        y += height
    presentation.save(output)


def wrong_fill(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    cell = target_table(presentation).table.cell(1, 1)
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor(255, 0, 255)
    presentation.save(output)


def left_aligned_values(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    for row in target_table(presentation).table.rows:
        row.cells[1].text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    presentation.save(output)


def overlapping_table(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    target_table(presentation).top = Inches(3.0)
    presentation.save(output)


def changed_title(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    shape_by_id(presentation, 6).text = "LOW-INCOME HOUSING TAX CREDIT"
    presentation.save(output)


def duplicate_target_table(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    slide = presentation.slides[0]
    target = target_table(presentation)
    slide.shapes.add_table(4, 2, target.left, target.top, target.width, target.height)
    presentation.save(output)


def removed_small_table(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    shape = shape_by_id(presentation, 11)
    shape._element.getparent().remove(shape._element)
    presentation.save(output)


def flattened_slide(golden: Path, slide_image: Path, output: Path) -> None:
    presentation = Presentation(golden)
    slide = presentation.slides[0]
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    slide.shapes.add_picture(str(slide_image), 0, 0, width=presentation.slide_width, height=presentation.slide_height)
    presentation.save(output)


def hardcoded_marker(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    box = presentation.slides[0].shapes.add_textbox(Inches(0.10), Inches(0.10), Inches(2.50), Inches(0.25))
    box.text = "solution/Golden"
    presentation.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--slide-image", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()

    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))
    valid_alternate(args.golden, output_path(args.outdir, "valid-alternate"))
    wrong_value(args.golden, output_path(args.outdir, "wrong-value"))
    swapped_columns(args.golden, output_path(args.outdir, "swapped-columns"))
    kept_source_textbox(args.golden, output_path(args.outdir, "kept-source-textbox"))
    fake_textbox_table(args.golden, output_path(args.outdir, "fake-textbox-table"))
    wrong_fill(args.golden, output_path(args.outdir, "wrong-fill"))
    left_aligned_values(args.golden, output_path(args.outdir, "left-aligned-values"))
    overlapping_table(args.golden, output_path(args.outdir, "overlapping-table"))
    changed_title(args.golden, output_path(args.outdir, "changed-title"))
    duplicate_target_table(args.golden, output_path(args.outdir, "duplicate-target-table"))
    removed_small_table(args.golden, output_path(args.outdir, "removed-small-table"))
    flattened_slide(args.golden, args.slide_image, output_path(args.outdir, "flattened-slide"))
    hardcoded_marker(args.golden, output_path(args.outdir, "hardcoded-marker"))
    extra = args.outdir / "extra-output-file"
    extra.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.golden, extra / OUTPUT_NAME)
    (extra / "debug.txt").write_text("debug\n", encoding="utf-8")


if __name__ == "__main__":
    main()
