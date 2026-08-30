#!/usr/bin/env python3
"""PPTArena Case 53 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import posixpath
import re
import statistics
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from lxml import etree
from PIL import Image, ImageChops, ImageDraw, ImageStat
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN


OUTPUT_NAME = "FillTablefromData_GroundTruthA.pptx"
INPUT_NAME = "FillTablefromData_TestA.pptx"
INPUT_SHA256 = "69eb80458f9c78ee867e27812af34f63d12483fb1e4b4a9dd3eb63c3cdc37f99"
TASK_ID = "cowork-evo/pptarena-053-text-data-to-editable-table"
RUBRIC_VERSION = "pptarena-case-053-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-053-deterministic-v0.1"
SOURCE_TEXT = (
    "Revenue 30% AMI Rents Annual Actualized Income $387,030 "
    "Annual Operating Expenses ($420,300) Annual Net Revenue ($33,270)"
)
EXPECTED = (
    ("Revenue", "30% AMI Rents"),
    ("Annual Actualized Income", "$387,030"),
    ("Annual Operating Expenses", "($420,300)"),
    ("Annual Net Revenue", "($33,270)"),
)
SUSPICIOUS_MARKERS = ["solution/golden", "judge-result", "artifact_dir", "pptarena-case-053", "/root/"]

SPECS = {
    "R001": ("PPTX 完整且仍为原尺寸 1 页文稿", 7, "hurdle"),
    "R002": ("4×2 表格的 8 个财务真值正确", 24, "hurdle"),
    "R003": ("原数据文本框已删除且无重复数据", 10, "hurdle"),
    "R004": ("目标是唯一原生可编辑 4×2 表格", 12, "hurdle"),
    "R005": ("目标表格结构、边框、填充与主题样式保持", 12, "hurdle"),
    "R006": ("标签/数值对齐和字体层级清晰", 9, "hurdle"),
    "R007": ("表格位置合理、在页内且无重叠", 8, "hurdle"),
    "R008": ("非目标对象、备注和媒体保持", 10, "hurdle"),
    "R009": ("整页可渲染、文字可见且输出卫生", 8, "positive"),
    "R010": ("无数据、可编辑性、保持或答案泄漏严重错误", -20, "penalty"),
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def safe(callback, default=""):
    try:
        return callback()
    except Exception:
        return default


def locate_source() -> Path | None:
    candidates = [
        Path("/app/input_files") / INPUT_NAME,
        Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME,
    ]
    return next((path for path in candidates if path.is_file()), None)


def media_hashes(path: Path) -> Counter[str]:
    with zipfile.ZipFile(path) as archive:
        return Counter(
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("ppt/media/")
        )


def package_checks(path: Path) -> tuple[bool, dict]:
    evidence = {"broken_member": None, "missing_targets": [], "macros": []}
    try:
        with zipfile.ZipFile(path) as archive:
            evidence["broken_member"] = archive.testzip()
            names = set(archive.namelist())
            evidence["macros"] = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
            for rels_name in sorted(name for name in names if name.endswith(".rels")):
                root = etree.fromstring(archive.read(rels_name))
                if rels_name == "_rels/.rels":
                    owner = ""
                else:
                    prefix, filename = rels_name.split("/_rels/", 1)
                    owner = posixpath.join(prefix, filename[:-5])
                base = posixpath.dirname(owner)
                for relation in root:
                    if relation.attrib.get("TargetMode") == "External":
                        continue
                    target = relation.attrib.get("Target", "")
                    resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                    if resolved not in names:
                        evidence["missing_targets"].append(f"{rels_name} -> {resolved}")
        passed = evidence["broken_member"] is None and not evidence["missing_targets"] and not evidence["macros"]
        return passed, evidence
    except Exception as exc:
        evidence["exception"] = str(exc)
        return False, evidence


def font_signature(font) -> tuple:
    color = font.color
    return (
        font.name,
        round(font.size.pt, 3) if font.size else None,
        font.bold,
        font.italic,
        font.underline,
        safe(lambda: str(font.language_id)),
        safe(lambda: str(color.type)),
        safe(lambda: str(color.rgb)),
        safe(lambda: str(color.theme_color)),
        safe(lambda: round(color.brightness, 5)),
    )


def paragraphs_signature(text_frame) -> tuple:
    return tuple(
        (
            paragraph.text,
            paragraph.level,
            str(paragraph.alignment),
            str(paragraph.line_spacing),
            str(paragraph.space_before),
            str(paragraph.space_after),
            tuple((run.text, font_signature(run.font)) for run in paragraph.runs),
        )
        for paragraph in text_frame.paragraphs
    )


def target_style_signature(shape) -> tuple:
    return (
        tuple(row.height for row in shape.table.rows),
        tuple(column.width for column in shape.table.columns),
        etree.tostring(shape._element.graphic.graphicData.tbl.tblPr),
        tuple(
            etree.tostring(cell._tc.get_or_add_tcPr())
            for row in shape.table.rows
            for cell in row.cells
        ),
    )


def table_content_signature(shape) -> tuple:
    return (
        target_style_signature(shape),
        tuple(
            tuple((cell.text, paragraphs_signature(cell.text_frame)) for cell in row.cells)
            for row in shape.table.rows
        ),
    )


def shape_signature(shape) -> tuple:
    base = (
        str(shape.shape_type),
        int(shape.left),
        int(shape.top),
        int(shape.width),
        int(shape.height),
        safe(lambda: int(shape.rotation), 0),
        str(shape.placeholder_format.type) if shape.is_placeholder else "",
    )
    text = paragraphs_signature(shape.text_frame) if getattr(shape, "has_text_frame", False) else ()
    table = table_content_signature(shape) if getattr(shape, "has_table", False) else ()
    picture = ()
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        picture = (
            hashlib.sha256(shape.image.blob).hexdigest(),
            round(shape.crop_left, 6),
            round(shape.crop_right, 6),
            round(shape.crop_top, 6),
            round(shape.crop_bottom, 6),
        )
    return base, text, table, picture


def shape_map(presentation: Presentation) -> dict[int, object]:
    return {shape.shape_id: shape for shape in presentation.slides[0].shapes}


def find_target_tables(presentation: Presentation) -> list:
    return [
        shape
        for shape in presentation.slides[0].shapes
        if shape.has_table and len(shape.table.rows) == 4 and len(shape.table.columns) == 2
    ]


def notes_signature(slide) -> tuple[str, ...]:
    try:
        return tuple(shape.text for shape in slide.notes_slide.shapes if getattr(shape, "has_text_frame", False))
    except Exception:
        return ()


def run_sizes(cell) -> list[float]:
    return [run.font.size.pt for paragraph in cell.text_frame.paragraphs for run in paragraph.runs if run.text and run.font.size]


def direct_bold(cell) -> list[bool | None]:
    return [run.font.bold for paragraph in cell.text_frame.paragraphs for run in paragraph.runs if run.text]


def last_value_color_ok(cell) -> bool:
    runs = [run for paragraph in cell.text_frame.paragraphs for run in paragraph.runs if run.text]
    for run in runs:
        theme = safe(lambda: str(run.font.color.theme_color)).lower()
        rgb = safe(lambda: str(run.font.color.rgb))
        if "accent_5" in theme or "accent 5" in theme:
            return True
        if re.fullmatch(r"[0-9A-Fa-f]{6}", rgb):
            red, green, blue = int(rgb[:2], 16), int(rgb[2:4], 16), int(rgb[4:], 16)
            if red >= 140 and red > green * 1.15 and red > blue * 1.15:
                return True
    return False


def intersect_area(first, second) -> int:
    left = max(first.left, second.left)
    top = max(first.top, second.top)
    right = min(first.left + first.width, second.left + second.width)
    bottom = min(first.top + first.height, second.top + second.height)
    return max(0, right - left) * max(0, bottom - top)


def render(path: Path, outdir: Path) -> tuple[list[Path], int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "profile"
    profile.mkdir()
    completed = subprocess.run(
        ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(path)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=300,
        check=False,
    )
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return [], 0, completed.stdout[-500:]
    info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=30, check=False).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    pages = int(match.group(1)) if match else 0
    prefix = outdir / "slide"
    raster = subprocess.run(["pdftoppm", "-png", "-r", "100", str(pdf), str(prefix)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300, check=False)
    images = sorted(outdir.glob("slide-*.png")) if raster.returncode == 0 else []
    return images, pages, f"soffice={completed.returncode}, raster={raster.returncode}"


def emu_box(shape, image: Image.Image, presentation: Presentation) -> tuple[int, int, int, int]:
    return (
        max(0, int(shape.left * image.width / presentation.slide_width)),
        max(0, int(shape.top * image.height / presentation.slide_height)),
        min(image.width, int(math.ceil((shape.left + shape.width) * image.width / presentation.slide_width))),
        min(image.height, int(math.ceil((shape.top + shape.height) * image.height / presentation.slide_height))),
    )


def render_metrics(source_image: Path, candidate_image: Path, source: Presentation, candidate: Presentation, target) -> tuple[float, float]:
    source_shapes = shape_map(source)
    with Image.open(source_image) as source_raw, Image.open(candidate_image) as candidate_raw:
        first = source_raw.convert("RGB")
        second = candidate_raw.convert("RGB")
        if first.size != second.size:
            return 0.0, 0.0
        mask = Image.new("L", first.size, 255)
        draw = ImageDraw.Draw(mask)
        for shape, presentation in ((source_shapes[9], source), (source_shapes[16], source), (target, candidate)):
            draw.rectangle(emu_box(shape, first, presentation), fill=0)
        difference = ImageChops.difference(first, second)
        difference.paste((0, 0, 0), mask=ImageChops.invert(mask))
        similarity = max(0.0, 1.0 - statistics.mean(ImageStat.Stat(difference).rms) / 255.0)
        crop = second.crop(emu_box(target, second, candidate)).convert("L")
        histogram = crop.histogram()
        dark_ratio = sum(histogram[:210]) / max(1, crop.width * crop.height)
        return similarity, dark_ratio


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {
        "criterion_id": identifier,
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
        "raw": {"description": description, "weight": weight, "type": kind, "passed": bool(passed)},
    }


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [item["criterion_id"] for item in criteria if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0 and item["raw"]["passed"])
    penalty = sum(abs(item["raw"]["weight"]) for item in criteria if item["raw"]["weight"] < 0 and not item["raw"]["passed"])
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [sha256(path)] if path.is_file() else [],
        "criterion_results": criteria,
        "reward": reward,
        "passed": reward >= 0.8 and not hurdles,
        "hurdle_failures": hurdles,
    }


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        return make_result(path, [result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配") for identifier in SPECS])
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        return make_result(path, [result_item(identifier, False, f"PPTX 打开失败：{exc}") for identifier in SPECS])

    criteria = []
    package_passed, package_evidence = package_checks(path)
    size_passed = len(source.slides) == len(candidate.slides) == 1 and (source.slide_width, source.slide_height) == (candidate.slide_width, candidate.slide_height)
    criteria.append(result_item("R001", package_passed and size_passed, f"页数={len(candidate.slides)}，尺寸一致={size_passed}，包={json.dumps(package_evidence, ensure_ascii=False)}"))

    targets = find_target_tables(candidate) if len(candidate.slides) == 1 else []
    target = targets[0] if len(targets) == 1 else None
    actual = tuple(tuple(cell.text for cell in row.cells) for row in target.table.rows) if target is not None else ()
    truth_passed = len(targets) == 1 and actual == EXPECTED
    criteria.append(result_item("R002", truth_passed, f"4x2 表格数={len(targets)}，单元格={actual!r}"))

    all_texts = [shape.text for shape in candidate.slides[0].shapes if getattr(shape, "has_text_frame", False)] if len(candidate.slides) == 1 else []
    source_box_present = any(text == SOURCE_TEXT for text in all_texts)
    fragments_outside_target = []
    needles = ["30% AMI Rents", "$387,030", "($420,300)", "($33,270)"]
    for shape in candidate.slides[0].shapes if len(candidate.slides) == 1 else []:
        if shape is target or shape.has_table or not getattr(shape, "has_text_frame", False):
            continue
        for needle in needles:
            if needle in shape.text:
                fragments_outside_target.append((shape.shape_id, needle))
    deletion_passed = not source_box_present and not fragments_outside_target
    criteria.append(result_item("R003", deletion_passed, f"原文本框={source_box_present}，表外数据={fragments_outside_target}"))

    all_tables = [shape for shape in candidate.slides[0].shapes if shape.has_table] if len(candidate.slides) == 1 else []
    full_slide_pictures = []
    if len(candidate.slides) == 1:
        page_area = candidate.slide_width * candidate.slide_height
        full_slide_pictures = [shape.name for shape in candidate.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.width * shape.height >= page_area * 0.80]
    editable_passed = target is not None and len(all_tables) == 3 and len(candidate.slides[0].shapes) == 8 and not full_slide_pictures
    criteria.append(result_item("R004", editable_passed, f"原生表格数={len(all_tables)}，页面对象={len(candidate.slides[0].shapes) if len(candidate.slides) == 1 else -1}，整页图={full_slide_pictures}"))

    source_target = shape_map(source)[16]
    style_passed = target is not None and target_style_signature(target) == target_style_signature(source_target)
    criteria.append(result_item("R005", style_passed, f"行高/列宽/表格主题/tcPr 一致={style_passed}"))

    hierarchy_passed = False
    hierarchy_evidence = {}
    if target is not None and truth_passed:
        left_align = all(row.cells[0].text_frame.paragraphs[0].alignment in (PP_ALIGN.LEFT, None) for row in target.table.rows)
        right_align = all(row.cells[1].text_frame.paragraphs[0].alignment == PP_ALIGN.RIGHT for row in target.table.rows)
        sizes = [[run_sizes(cell) for cell in row.cells] for row in target.table.rows]
        sizes_ok = all(sizes[i][j] and min(sizes[i][j]) >= (22 if (i, j) == (0, 0) else 14) and max(sizes[i][j]) <= 32 for i in range(4) for j in range(2))
        header_bold = True in direct_bold(target.table.cell(0, 0))
        result_bold = True in direct_bold(target.table.cell(3, 1))
        result_color = last_value_color_ok(target.table.cell(3, 1))
        hierarchy_passed = left_align and right_align and sizes_ok and header_bold and result_bold and result_color
        hierarchy_evidence = {"left_align": left_align, "right_align": right_align, "sizes": sizes, "header_bold": header_bold, "result_bold": result_bold, "result_color": result_color}
    criteria.append(result_item("R006", hierarchy_passed, json.dumps(hierarchy_evidence, ensure_ascii=False)))

    layout_passed = False
    layout_evidence = {}
    if target is not None and len(candidate.slides) == 1:
        source_shapes = shape_map(source)
        top_tables = [shape for shape in candidate.slides[0].shapes if shape.has_table and shape.shape_id != target.shape_id]
        body = shape_map(candidate).get(7)
        footer = shape_map(candidate).get(4)
        geometry_close = abs(target.left - source_target.left) <= 228_600 and abs(target.width - source_target.width) <= 91_440 and abs(target.height - source_target.height) <= 91_440
        top_range = 3_350_000 <= target.top <= 4_260_000
        inside = target.left >= 0 and target.top >= 0 and target.left + target.width <= candidate.slide_width and target.top + target.height <= candidate.slide_height
        overlaps = [(shape.shape_id, intersect_area(target, shape)) for shape in [*top_tables, body, footer] if shape is not None]
        no_overlap = all(area == 0 for _, area in overlaps)
        layout_passed = geometry_close and top_range and inside and no_overlap
        layout_evidence = {"left": target.left, "top": target.top, "width": target.width, "height": target.height, "geometry_close": geometry_close, "top_range": top_range, "inside": inside, "overlaps": overlaps}
    criteria.append(result_item("R007", layout_passed, json.dumps(layout_evidence, ensure_ascii=False)))

    preservation_passed = False
    preservation_evidence = {}
    if len(candidate.slides) == 1:
        source_shapes = shape_map(source)
        candidate_shapes = shape_map(candidate)
        required_ids = sorted(shape_id for shape_id in source_shapes if shape_id not in {9, 16})
        present = all(shape_id in candidate_shapes for shape_id in required_ids)
        signatures = present and all(shape_signature(candidate_shapes[shape_id]) == shape_signature(source_shapes[shape_id]) for shape_id in required_ids)
        notes_ok = notes_signature(candidate.slides[0]) == notes_signature(source.slides[0])
        media_ok = media_hashes(path) == media_hashes(source_path)
        preservation_passed = present and signatures and notes_ok and media_ok
        preservation_evidence = {"required_ids": required_ids, "present": present, "signatures": signatures, "notes": notes_ok, "media": media_ok}
    criteria.append(result_item("R008", preservation_passed, json.dumps(preservation_evidence, ensure_ascii=False)))

    markers = [marker for marker in SUSPICIOUS_MARKERS if any(marker.lower() in text.lower() for text in all_texts)]
    artifact_files = sorted(str(item.relative_to(artifact_dir)) for item in artifact_dir.rglob("*") if item.is_file())
    render_passed = False
    render_evidence = {"skipped": "前置条件未通过"}
    prerequisites = target is not None and all(next(item for item in criteria if item["criterion_id"] == identifier)["raw"]["passed"] for identifier in ("R001", "R002", "R003", "R004", "R005", "R006", "R007", "R008"))
    if prerequisites:
        with tempfile.TemporaryDirectory(prefix="pptarena053-judge-") as temporary:
            root = Path(temporary)
            source_images, source_pages, source_note = render(source_path, root / "source")
            candidate_images, candidate_pages, candidate_note = render(path, root / "candidate")
            similarity, dark_ratio = (0.0, 0.0)
            if len(source_images) == len(candidate_images) == 1:
                similarity, dark_ratio = render_metrics(source_images[0], candidate_images[0], source, candidate, target)
            render_passed = source_pages == candidate_pages == 1 and similarity >= 0.999 and dark_ratio >= 0.035
            render_evidence = {"source": source_note, "candidate": candidate_note, "source_pages": source_pages, "candidate_pages": candidate_pages, "non_target_similarity": round(similarity, 6), "target_dark_ratio": round(dark_ratio, 6)}
    hygiene_passed = render_passed and artifact_files == [OUTPUT_NAME] and not markers
    render_evidence["outputs"] = artifact_files
    render_evidence["markers"] = markers
    criteria.append(result_item("R009", hygiene_passed, json.dumps(render_evidence, ensure_ascii=False)))

    severe = not all(next(item for item in criteria if item["criterion_id"] == identifier)["raw"]["passed"] for identifier in ("R001", "R002", "R003", "R004", "R005", "R006", "R007", "R008")) or bool(markers) or not render_passed
    criteria.append(result_item("R010", not severe, f"严重错误={severe}"))
    order = {identifier: index for index, identifier in enumerate(SPECS)}
    criteria.sort(key=lambda item: order[item["criterion_id"]])
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    output = args.artifact_dir / OUTPUT_NAME
    if output.is_file():
        evaluation = evaluate(output, args.artifact_dir)
    else:
        evaluation = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [], "criterion_results": [], "reward": 0.0, "passed": False, "hurdle_failures": ["missing_output"]}
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
