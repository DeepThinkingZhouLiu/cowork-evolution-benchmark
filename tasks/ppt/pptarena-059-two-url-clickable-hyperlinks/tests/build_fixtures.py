#!/usr/bin/env python3
"""只为 PPTArena Case 59 构造可重放的合法变体、边界例和负例。"""

from __future__ import annotations

import argparse
import copy
import shutil
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Inches


OUTPUT_NAME = "ConvertHyperLink_GroundTruth.pptx"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
HYPERLINK_REL_TYPE = f"{R_NS}/hyperlink"
FULL_SENTENCE = "Visit https://www.berkeley.edu/ and https://illinois.edu/ for more details"


def qn(namespace: str, local: str) -> str:
    return f"{{{namespace}}}{local}"


def output_path(root: Path, case: str) -> Path:
    target = root / case / OUTPUT_NAME
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def serialize(root) -> bytes:
    return etree.tostring(
        root,
        encoding="UTF-8",
        xml_declaration=True,
        standalone=True,
    )


def rewrite_archive(source: Path, output: Path, transform) -> None:
    with zipfile.ZipFile(source) as archive_in, zipfile.ZipFile(output, "w") as archive_out:
        for info in archive_in.infolist():
            archive_out.writestr(info, transform(info.filename, archive_in.read(info.filename)))


def target_paragraph(root):
    matches = []
    for shape in root.xpath(".//p:sp", namespaces={"p": P_NS}):
        visible = "".join(shape.xpath(".//a:t/text()", namespaces={"a": A_NS}))
        if visible == FULL_SENTENCE:
            matches.append(shape)
    if len(matches) != 1:
        raise ValueError(f"目标文本框数量不是 1：{len(matches)}")
    paragraphs = matches[0].xpath("./p:txBody/a:p", namespaces={"p": P_NS, "a": A_NS})
    if len(paragraphs) != 1:
        raise ValueError("目标段落数不是 1")
    return paragraphs[0]


def replace_runs(
    root,
    parts: list[tuple[str, str | None]],
    *,
    explicit_style: bool = False,
    url_size: int | None = None,
) -> None:
    paragraph = target_paragraph(root)
    existing = paragraph.xpath("./a:r", namespaces={"a": A_NS})
    if not existing:
        raise ValueError("冻结输入无目标 run")
    template_properties = existing[0].find(qn(A_NS, "rPr"))
    insertion = paragraph.index(existing[0])
    for run in existing:
        paragraph.remove(run)
    for offset, (text, relationship_id) in enumerate(parts):
        run = etree.Element(qn(A_NS, "r"))
        properties = copy.deepcopy(template_properties)
        if properties is None:
            properties = etree.Element(qn(A_NS, "rPr"))
        if relationship_id and url_size is not None:
            properties.set("sz", str(url_size))
        if explicit_style and (relationship_id or text in {"https://www.berkeley.edu/", "https://illinois.edu/"}):
            properties.set("u", "sng")
            solid = etree.SubElement(properties, qn(A_NS, "solidFill"))
            etree.SubElement(solid, qn(A_NS, "srgbClr"), val="0563C1")
        if relationship_id:
            etree.SubElement(
                properties,
                qn(A_NS, "hlinkClick"),
                {qn(R_NS, "id"): relationship_id},
            )
        run.append(properties)
        node = etree.SubElement(run, qn(A_NS, "t"))
        if text.startswith(" ") or text.endswith(" "):
            node.set(qn("http://www.w3.org/XML/1998/namespace", "space"), "preserve")
        node.text = text
        paragraph.insert(insertion + offset, run)


def add_relationships(root, relationships: list[tuple[str, str, bool]]) -> None:
    for relationship_id, target, external in relationships:
        attributes = {
            "Id": relationship_id,
            "Type": HYPERLINK_REL_TYPE,
            "Target": target,
        }
        if external:
            attributes["TargetMode"] = "External"
        etree.SubElement(root, qn(PKG_REL_NS, "Relationship"), **attributes)


def hyperlink_variant(
    source: Path,
    output: Path,
    *,
    parts: list[tuple[str, str | None]],
    relationships: list[tuple[str, str, bool]],
    explicit_style: bool = False,
    url_size: int | None = None,
) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name == "ppt/slides/slide1.xml":
            root = etree.fromstring(data)
            replace_runs(
                root,
                parts,
                explicit_style=explicit_style,
                url_size=url_size,
            )
            return serialize(root)
        if name == "ppt/slides/_rels/slide1.xml.rels":
            root = etree.fromstring(data)
            add_relationships(root, relationships)
            return serialize(root)
        return data

    rewrite_archive(source, output, transform)


def standard_parts(first_id: str | None = "rId6", second_id: str | None = "rId7"):
    return [
        ("Visit ", None),
        ("https://www.berkeley.edu/", first_id),
        (" and ", None),
        ("https://illinois.edu/", second_id),
        (" for more details", None),
    ]


def changed_text(golden: Path, output: Path) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name != "ppt/slides/slide1.xml":
            return data
        root = etree.fromstring(data)
        paragraph = target_paragraph(root)
        nodes = paragraph.xpath("./a:r/a:t", namespaces={"a": A_NS})
        nodes[1].text = "https://www.berkeley.com/"
        return serialize(root)

    rewrite_archive(golden, output, transform)


def changed_title(golden: Path, output: Path) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name != "ppt/slides/slide1.xml":
            return data
        root = etree.fromstring(data)
        nodes = root.xpath(".//a:t", namespaces={"a": A_NS})
        title = next(node for node in nodes if node.text and node.text.startswith("WRITE AN ORIGINAL"))
        title.text = f"{title.text} X"
        return serialize(root)

    rewrite_archive(golden, output, transform)


def moved_background(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    background = presentation.slides[0].shapes[0]
    background.left += Inches(0.5)
    presentation.save(output)


def hardcoded_marker(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    box = presentation.slides[0].shapes.add_textbox(
        Inches(0.1), Inches(0.1), Inches(2.5), Inches(0.25)
    )
    box.text = "solution/Golden"
    presentation.save(output)


def extra_slide(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    presentation.slides.add_slide(presentation.slide_layouts[0])
    presentation.save(output)


def changed_image(golden: Path, output: Path) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name != "ppt/media/image1.png":
            return data
        with zipfile.ZipFile(golden) as archive:
            return archive.read("ppt/media/image2.png")

    rewrite_archive(golden, output, transform)


def repacked_content_types(golden: Path, output: Path) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name == "[Content_Types].xml":
            return data.replace(b"</Types>", b"  \n</Types>")
        return data

    rewrite_archive(golden, output, transform)


def shape_level_link(source: Path, output: Path) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name == "ppt/slides/slide1.xml":
            root = etree.fromstring(data)
            paragraph = target_paragraph(root)
            shape = paragraph.xpath("ancestor::p:sp[1]", namespaces={"p": P_NS})[0]
            properties = shape.xpath("./p:nvSpPr/p:cNvPr", namespaces={"p": P_NS})[0]
            etree.SubElement(properties, qn(A_NS, "hlinkClick"), {qn(R_NS, "id"): "rId6"})
            return serialize(root)
        if name == "ppt/slides/_rels/slide1.xml.rels":
            root = etree.fromstring(data)
            add_relationships(root, [("rId6", "https://www.berkeley.edu/", True)])
            return serialize(root)
        return data

    rewrite_archive(source, output, transform)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--upstream-unfit", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()

    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))
    shutil.copy2(args.upstream_unfit, output_path(args.outdir, "upstream-unfit"))

    hyperlink_variant(
        args.source,
        output_path(args.outdir, "valid-explicit-blue"),
        parts=standard_parts("rId20", "rId21"),
        relationships=[
            ("rId20", "https://www.berkeley.edu/", True),
            ("rId21", "https://illinois.edu/", True),
        ],
        explicit_style=True,
    )
    hyperlink_variant(
        args.source,
        output_path(args.outdir, "valid-split-url-runs"),
        parts=[
            ("Visit ", None),
            ("https://www.", "rId20"),
            ("berkeley.edu/", "rId20"),
            (" and ", None),
            ("https://", "rId21"),
            ("illinois.edu/", "rId21"),
            (" for more details", None),
        ],
        relationships=[
            ("rId20", "https://www.berkeley.edu/", True),
            ("rId21", "https://illinois.edu/", True),
        ],
    )
    repacked_content_types(args.golden, output_path(args.outdir, "repacked-content-types"))

    hyperlink_variant(
        args.source,
        output_path(args.outdir, "style-only"),
        parts=standard_parts(None, None),
        relationships=[],
        explicit_style=True,
    )
    hyperlink_variant(
        args.source,
        output_path(args.outdir, "only-berkeley"),
        parts=standard_parts("rId6", None),
        relationships=[("rId6", "https://www.berkeley.edu/", True)],
    )
    hyperlink_variant(
        args.source,
        output_path(args.outdir, "only-illinois"),
        parts=standard_parts(None, "rId7"),
        relationships=[("rId7", "https://illinois.edu/", True)],
    )
    hyperlink_variant(
        args.source,
        output_path(args.outdir, "swapped-targets"),
        parts=standard_parts("rId6", "rId7"),
        relationships=[
            ("rId6", "https://illinois.edu/", True),
            ("rId7", "https://www.berkeley.edu/", True),
        ],
    )
    hyperlink_variant(
        args.source,
        output_path(args.outdir, "whole-sentence-linked"),
        parts=[(FULL_SENTENCE, "rId6")],
        relationships=[("rId6", "https://www.berkeley.edu/", True)],
    )
    hyperlink_variant(
        args.source,
        output_path(args.outdir, "surrounding-linked"),
        parts=[
            ("Visit ", "rId6"),
            ("https://www.berkeley.edu/", "rId6"),
            (" and ", None),
            ("https://illinois.edu/", "rId7"),
            (" for more details", None),
        ],
        relationships=[
            ("rId6", "https://www.berkeley.edu/", True),
            ("rId7", "https://illinois.edu/", True),
        ],
    )
    hyperlink_variant(
        args.source,
        output_path(args.outdir, "wrong-http-scheme"),
        parts=standard_parts("rId6", "rId7"),
        relationships=[
            ("rId6", "http://www.berkeley.edu/", True),
            ("rId7", "http://illinois.edu/", True),
        ],
    )
    hyperlink_variant(
        args.source,
        output_path(args.outdir, "internal-target-mode"),
        parts=standard_parts("rId6", "rId7"),
        relationships=[
            ("rId6", "https://www.berkeley.edu/", False),
            ("rId7", "https://illinois.edu/", False),
        ],
    )
    shape_level_link(args.source, output_path(args.outdir, "shape-level-link"))
    hyperlink_variant(
        args.source,
        output_path(args.outdir, "wrong-font-size"),
        parts=standard_parts("rId6", "rId7"),
        relationships=[
            ("rId6", "https://www.berkeley.edu/", True),
            ("rId7", "https://illinois.edu/", True),
        ],
        url_size=1400,
    )
    changed_text(args.golden, output_path(args.outdir, "changed-url-text"))
    changed_title(args.golden, output_path(args.outdir, "changed-title"))
    moved_background(args.golden, output_path(args.outdir, "moved-background"))
    changed_image(args.golden, output_path(args.outdir, "changed-image"))
    hardcoded_marker(args.golden, output_path(args.outdir, "hardcoded-marker"))
    extra_slide(args.golden, output_path(args.outdir, "extra-slide"))


if __name__ == "__main__":
    main()
