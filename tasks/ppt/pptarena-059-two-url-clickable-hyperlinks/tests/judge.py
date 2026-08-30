#!/usr/bin/env python3
"""PPTArena Case 59 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from lxml import etree
from PIL import Image, ImageChops
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "ConvertHyperLink_GroundTruth.pptx"
INPUT_NAME = "ConvertHyperLink_Test.pptx"
INPUT_SHA256 = "33132900f1422268aec0a20ba9e6ba0c2be1e6b9a4baeb9f0887c371506792f5"
TASK_ID = "cowork-evo/pptarena-059-two-url-clickable-hyperlinks"
RUBRIC_VERSION = "pptarena-case-059-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-059-deterministic-v0.1"
FULL_SENTENCE = "Visit https://www.berkeley.edu/ and https://illinois.edu/ for more details"
URL_TARGETS = {
    "https://www.berkeley.edu/": "https://www.berkeley.edu/",
    "https://illinois.edu/": "https://illinois.edu/",
}
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
HYPERLINK_REL_TYPE = f"{R_NS}/hyperlink"
EMU_PER_INCH = 914_400
TARGET_SHAPE_ID = 918
SUSPICIOUS_MARKERS = [
    "solution/golden",
    "/root/",
    "judge-result",
    "artifact_dir",
    "pptarena-case-059",
]

SPECS = {
    "R001": ("指定 PPTX 可打开、页面尺寸不变且总页数为 1", 6, "hurdle"),
    "R002": ("两个 URL 的全部字符分别具有正确的外部 hyperlink", 22, "hurdle"),
    "R003": ("只有两个 URL 被链接，周围文字和其他对象无链接/动作", 12, "hurdle"),
    "R004": ("整句文字、顺序、12pt 字号和字体属性保持", 15, "hurdle"),
    "R005": ("全部原页面对象的内容、几何、裁剪、顺序、版式和样式保持", 12, "hurdle"),
    "R006": ("恰好存在两个正确 TargetMode=External 的 slide1 hyperlink 关系", 10, "hurdle"),
    "R007": ("保持 15 个原生页面对象、3 个原媒体和原包成员集", 8, "hurdle"),
    "R008": ("页面可渲染，除 URL 默认样式区外与冻结输入高度一致", 8, "hurdle"),
    "R009": ("输出目录干净且不含答案路径或未授权文字", 5, "positive"),
    "R010": ("规范化掉合法 hyperlink 后内容 XML 和关键包部件字节级保持", 2, "positive"),
    "R011": ("URL 缺链/错链/伪样式或原页面内容、布局、可编辑性被破坏", -20, "penalty"),
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def qn(namespace: str, local: str) -> str:
    return f"{{{namespace}}}{local}"


def locate_source() -> Path | None:
    candidates = [
        Path("/app/input_files") / INPUT_NAME,
        Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME,
    ]
    return next((path for path in candidates if path.is_file()), None)


def normalize_visible(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def rounded_geometry(shape) -> tuple:
    return tuple(
        round(int(value) / 1000) * 1000
        for value in (shape.left, shape.top, shape.width, shape.height)
    )


def color_signature(color) -> tuple:
    try:
        color_type = int(color.type) if color.type is not None else None
    except Exception:
        color_type = str(getattr(color, "type", None))
    try:
        rgb = str(color.rgb) if color.rgb is not None else None
    except Exception:
        rgb = None
    try:
        theme = int(color.theme_color) if color.theme_color is not None else None
    except Exception:
        theme = None
    try:
        brightness = round(color.brightness, 5)
    except Exception:
        brightness = None
    return color_type, rgb, theme, brightness


def font_signature(font, run=None, *, ignore_hyperlink_style: bool = False) -> tuple:
    try:
        language_value = font.language_id
        language = int(language_value) if language_value is not None else None
    except Exception:
        language = None
    try:
        hyperlink = run.hyperlink.address if run is not None else None
    except Exception:
        hyperlink = None
    color = None if ignore_hyperlink_style and hyperlink else color_signature(font.color)
    underline = None if ignore_hyperlink_style and hyperlink else bool(font.underline)
    return (
        font.name,
        round(font.size.pt, 4) if font.size else None,
        bool(font.bold),
        bool(font.italic),
        underline,
        language,
        color,
    )


def text_signature(shape) -> tuple:
    text_frame = shape.text_frame
    paragraphs = []
    for paragraph in text_frame.paragraphs:
        runs = tuple(
            (run.text, font_signature(run.font, run))
            for run in paragraph.runs
        )
        paragraphs.append(
            (
                paragraph.text,
                paragraph.level,
                str(paragraph.alignment),
                str(paragraph.line_spacing),
                str(paragraph.space_before),
                str(paragraph.space_after),
                runs,
            )
        )
    return (
        (
            text_frame.margin_left,
            text_frame.margin_right,
            text_frame.margin_top,
            text_frame.margin_bottom,
            str(text_frame.vertical_anchor),
            text_frame.word_wrap,
            str(text_frame.auto_size),
        ),
        tuple(paragraphs),
    )


def picture_signature(shape) -> tuple:
    try:
        blob = hashlib.sha256(shape.image.blob).hexdigest()
    except Exception as exc:
        blob = f"unreadable:{type(exc).__name__}"
    return (
        blob,
        round(float(shape.crop_left), 6),
        round(float(shape.crop_right), 6),
        round(float(shape.crop_top), 6),
        round(float(shape.crop_bottom), 6),
    )


def shape_signature(shape) -> tuple:
    text = text_signature(shape) if getattr(shape, "has_text_frame", False) else ()
    picture = picture_signature(shape) if shape.shape_type == MSO_SHAPE_TYPE.PICTURE else ()
    children = (
        tuple(shape_signature(child) for child in shape.shapes)
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP
        else ()
    )
    return (
        int(shape.shape_id),
        shape.name,
        str(shape.shape_type),
        rounded_geometry(shape),
        round(float(shape.rotation or 0), 4),
        text,
        picture,
        children,
    )


def target_shape(presentation: Presentation):
    matches = [
        shape
        for shape in presentation.slides[0].shapes
        if getattr(shape, "has_text_frame", False) and shape.text == FULL_SENTENCE
    ] if len(presentation.slides) else []
    return matches[0] if len(matches) == 1 else None


def target_skeleton(shape) -> tuple:
    if shape is None:
        return ()
    text_frame = shape.text_frame
    paragraph_skeleton = tuple(
        (
            paragraph.level,
            str(paragraph.alignment),
            str(paragraph.line_spacing),
            str(paragraph.space_before),
            str(paragraph.space_after),
        )
        for paragraph in text_frame.paragraphs
    )
    return (
        int(shape.shape_id),
        shape.name,
        str(shape.shape_type),
        rounded_geometry(shape),
        round(float(shape.rotation or 0), 4),
        text_frame.margin_left,
        text_frame.margin_right,
        text_frame.margin_top,
        text_frame.margin_bottom,
        str(text_frame.vertical_anchor),
        text_frame.word_wrap,
        str(text_frame.auto_size),
        paragraph_skeleton,
    )


def notes_signature(presentation: Presentation) -> tuple:
    values = []
    for slide in presentation.slides:
        try:
            values.append(slide.notes_slide.notes_text_frame.text)
        except Exception:
            values.append("")
    return tuple(values)


def target_run_coverage(shape) -> tuple[bool, dict]:
    evidence = {"runs": [], "text": None, "coverage": [], "reason": None}
    if shape is None or len(shape.text_frame.paragraphs) != 1:
        evidence["reason"] = "目标文本框缺失或段落数不是 1"
        return False, evidence
    paragraph = shape.text_frame.paragraphs[0]
    position = 0
    coverage: list[str | None] = []
    style_passed = True
    for run in paragraph.runs:
        try:
            address = run.hyperlink.address
        except Exception:
            address = None
        text = run.text
        size = round(run.font.size.pt, 4) if run.font.size else None
        name = run.font.name
        bold = run.font.bold
        italic = run.font.italic
        underline = run.font.underline
        try:
            color = color_signature(run.font.color)
        except Exception:
            color = (None, None, None, None)
        in_url = address is not None
        name_ok = name in {None, "Aptos", "+mn-lt"}
        size_ok = size == 12.0
        emphasis_ok = bold in {None, False} and italic in {None, False}
        surrounding_style_ok = in_url or (
            underline in {None, False}
            and color in {
                (None, None, None, 0),
                (None, None, None, None),
            }
        )
        style_ok = name_ok and size_ok and emphasis_ok and surrounding_style_ok
        style_passed = style_passed and style_ok
        evidence["runs"].append(
            {
                "start": position,
                "end": position + len(text),
                "text": text,
                "address": address,
                "size_pt": size,
                "font_name": name,
                "bold": bold,
                "italic": italic,
                "underline": underline,
                "color": color,
                "style_ok": style_ok,
            }
        )
        coverage.extend([address] * len(text))
        position += len(text)
    visible = "".join(run.text for run in paragraph.runs)
    evidence["text"] = visible
    evidence["coverage"] = coverage
    if visible != FULL_SENTENCE or len(coverage) != len(FULL_SENTENCE):
        evidence["reason"] = "可见文字与冻结输入不一致"
        return False, {**evidence, "style_passed": style_passed}
    expected: list[str | None] = [None] * len(FULL_SENTENCE)
    for url, target in URL_TARGETS.items():
        start = FULL_SENTENCE.index(url)
        expected[start : start + len(url)] = [target] * len(url)
    coverage_passed = coverage == expected
    evidence["style_passed"] = style_passed
    evidence["coverage_passed"] = coverage_passed
    return coverage_passed, evidence


def package_ok(path: Path) -> tuple[bool, str]:
    try:
        with zipfile.ZipFile(path) as archive:
            broken = archive.testzip()
            names = set(archive.namelist())
            required = {
                "[Content_Types].xml",
                "ppt/presentation.xml",
                "ppt/slides/slide1.xml",
                "ppt/slides/_rels/slide1.xml.rels",
            }
            missing = sorted(required - names)
            macros = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
        return broken is None and not missing and not macros, (
            f"损坏成员={broken}，缺失={missing}，宏={macros}"
        )
    except Exception as exc:
        return False, f"ZIP 检查失败：{exc}"


def raw_slide(path: Path):
    with zipfile.ZipFile(path) as archive:
        return etree.fromstring(archive.read("ppt/slides/slide1.xml"))


def raw_rels(path: Path):
    with zipfile.ZipFile(path) as archive:
        return etree.fromstring(archive.read("ppt/slides/_rels/slide1.xml.rels"))


def relationship_inventory(path: Path) -> dict:
    with zipfile.ZipFile(path) as archive:
        all_hyperlinks = []
        for name in archive.namelist():
            if not name.endswith(".rels"):
                continue
            root = etree.fromstring(archive.read(name))
            for relationship in root:
                if relationship.get("Type") == HYPERLINK_REL_TYPE:
                    all_hyperlinks.append(
                        {
                            "part": name,
                            "id": relationship.get("Id"),
                            "target": relationship.get("Target"),
                            "target_mode": relationship.get("TargetMode"),
                        }
                    )
    root = raw_slide(path)
    clicks = []
    for node in root.xpath(".//a:hlinkClick | .//a:hlinkMouseOver", namespaces={"a": A_NS}):
        ancestor_shapes = node.xpath("ancestor::p:sp[1]", namespaces={"p": P_NS})
        shape_id = None
        if ancestor_shapes:
            ids = ancestor_shapes[0].xpath("./p:nvSpPr/p:cNvPr/@id", namespaces={"p": P_NS})
            shape_id = int(ids[0]) if ids and ids[0].isdigit() else None
        ancestor_runs = node.xpath("ancestor::a:r[1]", namespaces={"a": A_NS})
        run_text = "".join(ancestor_runs[0].xpath("./a:t/text()", namespaces={"a": A_NS})) if ancestor_runs else None
        clicks.append(
            {
                "kind": etree.QName(node).localname,
                "relationship_id": node.get(qn(R_NS, "id")),
                "shape_id": shape_id,
                "run_text": run_text,
                "inside_run_properties": bool(node.xpath("parent::a:rPr", namespaces={"a": A_NS})),
            }
        )
    target_relationships = [
        item
        for item in all_hyperlinks
        if item["part"] == "ppt/slides/_rels/slide1.xml.rels"
    ]
    relationship_map = {item["id"]: item for item in target_relationships}
    click_ids = [item["relationship_id"] for item in clicks]
    exact_relationships = (
        len(all_hyperlinks) == 2
        and len(target_relationships) == 2
        and Counter(item["target"] for item in target_relationships) == Counter(URL_TARGETS.values())
        and all(item["target_mode"] == "External" for item in target_relationships)
    )
    exact_clicks = (
        len(clicks) >= 2
        and all(
            item["kind"] == "hlinkClick"
            and item["shape_id"] == TARGET_SHAPE_ID
            and item["inside_run_properties"]
            and item["relationship_id"] in relationship_map
            and bool(item["run_text"])
            and item["run_text"] in relationship_map[item["relationship_id"]]["target"]
            for item in clicks
        )
        and set(click_ids) == set(relationship_map)
    )
    return {
        "all_relationships": all_hyperlinks,
        "clicks": clicks,
        "exact_relationships": exact_relationships,
        "exact_clicks": exact_clicks,
    }


def media_hashes(path: Path) -> Counter[str]:
    with zipfile.ZipFile(path) as archive:
        return Counter(
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("ppt/media/")
        )


def stable_member_names(path: Path) -> set[str]:
    with zipfile.ZipFile(path) as archive:
        return {name for name in archive.namelist() if not name.endswith("/")}


def critical_hashes(path: Path) -> dict[str, str]:
    excluded = {
        "ppt/slides/slide1.xml",
        "ppt/slides/_rels/slide1.xml.rels",
    }
    with zipfile.ZipFile(path) as archive:
        return {
            name: hashlib.sha256(archive.read(name)).hexdigest()
            for name in sorted(archive.namelist())
            if not name.endswith("/") and name not in excluded and not name.startswith("docProps/")
        }


def raw_target_shape(root):
    matches = []
    for shape in root.xpath(".//p:sp", namespaces={"p": P_NS}):
        visible = "".join(shape.xpath(".//a:t/text()", namespaces={"a": A_NS}))
        if visible == FULL_SENTENCE:
            matches.append(shape)
    return matches[0] if len(matches) == 1 else None


def normalized_slide_hash(path: Path, source_path: Path) -> str:
    candidate = raw_slide(path)
    source = raw_slide(source_path)
    candidate_target = raw_target_shape(candidate)
    source_target = raw_target_shape(source)
    if candidate_target is None or source_target is None:
        return ""
    candidate_target.getparent().replace(candidate_target, etree.fromstring(etree.tostring(source_target)))
    c_sld = candidate.find(qn(P_NS, "cSld"))
    return hashlib.sha256(etree.tostring(c_sld, method="c14n")).hexdigest()


def normalized_rels_hash(path: Path) -> str:
    root = raw_rels(path)
    for relationship in list(root):
        if relationship.get("Type") == HYPERLINK_REL_TYPE:
            root.remove(relationship)
    return hashlib.sha256(etree.tostring(root, method="c14n")).hexdigest()


def render_presentation(path: Path, outdir: Path) -> tuple[list[Path], int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "lo-profile"
    profile.mkdir()
    completed = subprocess.run(
        [
            "soffice",
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(outdir),
            str(path),
        ],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return [], 0, f"转 PDF 失败：{completed.stdout[-500:]}"
    info = subprocess.run(
        ["pdfinfo", str(pdf)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=30,
        check=False,
    ).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    pages = int(match.group(1)) if match else 0
    prefix = outdir / "slide"
    raster = subprocess.run(
        ["pdftoppm", "-jpeg", "-r", "120", str(pdf), str(prefix)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    images = sorted(outdir.glob("slide-*.jpg")) if raster.returncode == 0 else []
    return images, pages, (
        f"页数={pages}，渲染图={len(images)}，"
        f"soffice={completed.returncode}，pdftoppm={raster.returncode}"
    )


def outside_target_similarity(
    source_image: Image.Image,
    candidate_image: Image.Image,
    shape,
    slide_size: tuple[int, int],
) -> tuple[float, dict]:
    left = round(shape.left * source_image.width / slide_size[0])
    top = round(shape.top * source_image.height / slide_size[1])
    right = round((shape.left + shape.width) * source_image.width / slide_size[0])
    bottom = round((shape.top + shape.height) * source_image.height / slide_size[1])
    padding = max(4, round(source_image.height * 0.015))
    mask_box = (
        max(0, left - padding),
        max(0, top - padding),
        min(source_image.width, right + padding),
        min(source_image.height, bottom + padding),
    )
    source = source_image.convert("RGB")
    candidate = candidate_image.convert("RGB")
    if source.size != candidate.size:
        return 0.0, {"reason": "渲染尺寸不一致", "source": source.size, "candidate": candidate.size}
    difference = ImageChops.difference(source, candidate)
    mask = Image.new("L", source.size, 255)
    mask.paste(0, mask_box)
    pixels = mask.histogram()[255]
    histogram = difference.histogram(mask)
    squared = sum(count * ((index % 256) ** 2) for index, count in enumerate(histogram))
    rms = math.sqrt(squared / (pixels * 3)) if pixels else 255.0
    return max(0.0, 1.0 - rms / 255.0), {
        "mask_box": list(mask_box),
        "evaluated_pixels": pixels,
        "rms": round(rms, 6),
    }


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {
        "criterion_id": identifier,
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
        "raw": {
            "description": description,
            "weight": weight,
            "type": kind,
            "passed": bool(passed),
        },
    }


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [
        item["criterion_id"]
        for item in criteria
        if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]
    ]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(
        item["raw"]["weight"]
        for item in criteria
        if item["raw"]["weight"] > 0 and item["raw"]["passed"]
    )
    penalty = sum(
        abs(item["raw"]["weight"])
        for item in criteria
        if item["raw"]["weight"] < 0 and not item["raw"]["passed"]
    )
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [sha256(path)] if path.is_file() else [],
        "criterion_results": criteria,
        "reward": reward,
        "passed": reward >= 0.8 and not hurdles,
        "hurdle_failures": hurdles,
    }


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        criteria = [
            result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配")
            for identifier in SPECS
        ]
        return make_result(path, criteria)
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        criteria = [
            result_item(identifier, False, f"PPTX 打开失败：{exc}")
            for identifier in SPECS
        ]
        return make_result(path, criteria)

    criteria: list[dict] = []
    package_passed, package_evidence = package_ok(path)
    size_passed = (
        len(source.slides) == 1
        and len(candidate.slides) == 1
        and candidate.slide_width == source.slide_width
        and candidate.slide_height == source.slide_height
    )
    criteria.append(
        result_item(
            "R001",
            package_passed and size_passed,
            f"{package_evidence}，输入页数={len(source.slides)}，"
            f"候选页数={len(candidate.slides)}，页面="
            f"{candidate.slide_width / EMU_PER_INCH:.4f}×"
            f"{candidate.slide_height / EMU_PER_INCH:.4f}英寸",
        )
    )

    source_target = target_shape(source)
    candidate_target = target_shape(candidate)
    coverage_passed, coverage = target_run_coverage(candidate_target)
    criteria.append(
        result_item(
            "R002",
            coverage_passed,
            f"按字符链接覆盖={json.dumps(coverage, ensure_ascii=False)}",
        )
    )

    try:
        relationships = relationship_inventory(path)
    except Exception as exc:
        relationships = {
            "all_relationships": [],
            "clicks": [],
            "exact_relationships": False,
            "exact_clicks": False,
            "error": str(exc),
        }
    only_urls_linked = coverage_passed and bool(relationships.get("exact_clicks"))
    criteria.append(
        result_item(
            "R003",
            only_urls_linked,
            f"链接/动作节点={json.dumps(relationships.get('clicks', []), ensure_ascii=False)}",
        )
    )

    text_and_style_passed = (
        candidate_target is not None
        and candidate_target.text == FULL_SENTENCE
        and bool(coverage.get("style_passed"))
    )
    criteria.append(
        result_item(
            "R004",
            text_and_style_passed,
            f"目标文字完整={candidate_target is not None and candidate_target.text == FULL_SENTENCE}，"
            f"运行样式通过={coverage.get('style_passed', False)}",
        )
    )

    preservation_passed = False
    preservation_evidence = {"reason": "页数或目标文本框不正确"}
    if len(candidate.slides) == 1 and source_target is not None and candidate_target is not None:
        source_others = [
            shape_signature(shape)
            for shape in source.slides[0].shapes
            if shape.shape_id != TARGET_SHAPE_ID
        ]
        candidate_others = [
            shape_signature(shape)
            for shape in candidate.slides[0].shapes
            if shape.shape_id != TARGET_SHAPE_ID
        ]
        other_shapes_equal = source_others == candidate_others
        target_skeleton_equal = target_skeleton(source_target) == target_skeleton(candidate_target)
        try:
            source_layout = str(source.slides[0].slide_layout.part.partname)
            candidate_layout = str(candidate.slides[0].slide_layout.part.partname)
        except Exception:
            source_layout = candidate_layout = ""
        notes_equal = notes_signature(source) == notes_signature(candidate)
        preservation_passed = (
            len(source.slides[0].shapes) == 15
            and len(candidate.slides[0].shapes) == 15
            and other_shapes_equal
            and target_skeleton_equal
            and source_layout == candidate_layout
            and notes_equal
        )
        preservation_evidence = {
            "shape_count": [len(source.slides[0].shapes), len(candidate.slides[0].shapes)],
            "other_shapes_equal": other_shapes_equal,
            "target_skeleton_equal": target_skeleton_equal,
            "layouts": [source_layout, candidate_layout],
            "notes_equal": notes_equal,
        }
    criteria.append(
        result_item(
            "R005",
            preservation_passed,
            json.dumps(preservation_evidence, ensure_ascii=False),
        )
    )

    relationships_passed = bool(relationships.get("exact_relationships")) and bool(
        relationships.get("exact_clicks")
    )
    criteria.append(
        result_item(
            "R006",
            relationships_passed,
            f"关系清单={json.dumps(relationships, ensure_ascii=False)}",
        )
    )

    try:
        source_media = media_hashes(source_path)
        candidate_media = media_hashes(path)
        source_members = stable_member_names(source_path)
        candidate_members = stable_member_names(path)
    except Exception:
        source_media = candidate_media = Counter()
        source_members = candidate_members = set()
    editable_passed = (
        package_passed
        and len(candidate.slides) == 1
        and len(candidate.slides[0].shapes) == 15
        and source_media == candidate_media
        and sum(candidate_media.values()) == 3
        and source_members == candidate_members
    )
    criteria.append(
        result_item(
            "R007",
            editable_passed,
            f"对象数={len(candidate.slides[0].shapes) if len(candidate.slides) else 0}，"
            f"媒体哈希相等={source_media == candidate_media}，"
            f"媒体数={sum(source_media.values())}→{sum(candidate_media.values())}，"
            f"包成员集相等={source_members == candidate_members}",
        )
    )

    artifact_files = sorted(
        str(item.relative_to(artifact_dir))
        for item in artifact_dir.rglob("*")
        if item.is_file()
    )
    all_text = "\n".join(
        shape.text
        for slide in candidate.slides
        for shape in iter_shapes(slide.shapes)
        if getattr(shape, "has_text_frame", False)
    )
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker.lower() in all_text.lower()]
    hygiene_passed = artifact_files == [OUTPUT_NAME] and not markers
    criteria.append(
        result_item(
            "R009",
            hygiene_passed,
            f"输出文件={artifact_files}，可疑标记={markers}",
        )
    )

    try:
        exact_slide = normalized_slide_hash(path, source_path) == normalized_slide_hash(source_path, source_path)
        exact_rels = normalized_rels_hash(path) == normalized_rels_hash(source_path)
        critical_equal = critical_hashes(path) == critical_hashes(source_path)
    except Exception:
        exact_slide = exact_rels = critical_equal = False
    criteria.append(
        result_item(
            "R010",
            exact_slide and exact_rels and critical_equal,
            f"归一化幻灯片保持={exact_slide}，"
            f"归一化关系保持={exact_rels}，关键包部件保持={critical_equal}",
        )
    )

    pre_render_ready = (
        package_passed
        and size_passed
        and coverage_passed
        and only_urls_linked
        and text_and_style_passed
        and preservation_passed
        and relationships_passed
        and editable_passed
        and source_target is not None
        and candidate_target is not None
    )
    render_passed = False
    similarity = 0.0
    similarity_evidence = {"reason": "前置结构或保持检查失败"}
    render_note = "未执行渲染"
    if pre_render_ready:
        with tempfile.TemporaryDirectory(prefix="pptarena059-judge-") as temporary:
            root = Path(temporary)
            source_images, source_pages, source_note = render_presentation(source_path, root / "source")
            candidate_images, candidate_pages, candidate_note = render_presentation(path, root / "candidate")
            render_note = f"源[{source_note}]；候选[{candidate_note}]"
            if len(source_images) == 1 and len(candidate_images) == 1:
                with Image.open(source_images[0]) as source_image, Image.open(candidate_images[0]) as candidate_image:
                    similarity, similarity_evidence = outside_target_similarity(
                        source_image,
                        candidate_image,
                        source_target,
                        (source.slide_width, source.slide_height),
                    )
            render_passed = (
                source_pages == 1
                and candidate_pages == 1
                and len(source_images) == 1
                and len(candidate_images) == 1
                and similarity >= 0.9985
            )
    criteria.append(
        result_item(
            "R008",
            render_passed,
            f"URL 文本区外像素相似度={similarity:.6f}，"
            f"证据={json.dumps(similarity_evidence, ensure_ascii=False)}；{render_note}",
        )
    )

    severe = (
        not pre_render_ready
        or not render_passed
        or bool(markers)
    )
    criteria.append(result_item("R011", not severe, f"严重错误={severe}"))
    order = {identifier: index for index, identifier in enumerate(SPECS)}
    criteria.sort(key=lambda item: order[item["criterion_id"]])
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    output = args.artifact_dir / OUTPUT_NAME
    if output.is_file():
        evaluation = evaluate(output, args.artifact_dir)
    else:
        evaluation = {
            "task_id": TASK_ID,
            "rubric_version": RUBRIC_VERSION,
            "judge_version": JUDGE_VERSION,
            "artifact_sha256": [],
            "criterion_results": [],
            "reward": 0.0,
            "passed": False,
            "hurdle_failures": ["missing_output"],
        }
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(
        json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
