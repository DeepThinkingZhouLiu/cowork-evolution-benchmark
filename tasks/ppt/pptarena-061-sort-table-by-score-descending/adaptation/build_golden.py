#!/usr/bin/env python3
"""从 PPTArena Case 61 冻结输入构建只重排表格数据行的本地 Golden。"""

from __future__ import annotations

import argparse
import copy
import hashlib
import zipfile
from collections import Counter
from pathlib import Path

from lxml import etree
from pptx import Presentation


INPUT_SHA256 = "dd2ee0141319451ea8b33b20c06b49e2985932f302ba6de1349566a864f7995c"
SOURCE_ROWS = (
    ("Name", "Team", "Score"),
    ("Alice", "A", "82"),
    ("Bob", "B", "96"),
    ("Cara", "C", "74"),
    ("Diego", "A", "88"),
    ("Eve", "B", "91"),
)
ORDER = (0, 2, 5, 4, 1, 3)
EXPECTED_ROWS = tuple(SOURCE_ROWS[index] for index in ORDER)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def media_hashes(path: Path) -> Counter[str]:
    with zipfile.ZipFile(path) as archive:
        return Counter(hashlib.sha256(archive.read(name)).hexdigest() for name in archive.namelist() if name.startswith("ppt/media/"))


def shape_by_id(presentation: Presentation, shape_id: int):
    matches = [shape for shape in presentation.slides[0].shapes if shape.shape_id == shape_id]
    if len(matches) != 1:
        raise ValueError(f"无法唯一定位对象 {shape_id}")
    return matches[0]


def table_style_signature(shape) -> tuple:
    return (
        shape.left,
        shape.top,
        shape.width,
        shape.height,
        tuple(row.height for row in shape.table.rows),
        tuple(column.width for column in shape.table.columns),
        etree.tostring(shape._element.graphic.graphicData.tbl.tblPr),
        tuple(etree.tostring(cell._tc.get_or_add_tcPr()) for row in shape.table.rows for cell in row.cells),
    )


def picture_signature(shape) -> tuple:
    return (
        shape.left,
        shape.top,
        shape.width,
        shape.height,
        round(shape.crop_left, 8),
        round(shape.crop_right, 8),
        round(shape.crop_top, 8),
        round(shape.crop_bottom, 8),
        hashlib.sha256(shape.image.blob).hexdigest(),
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if not args.source.is_file() or sha256(args.source) != INPUT_SHA256:
        raise ValueError("冻结输入缺失或 SHA256 不匹配")

    presentation = Presentation(args.source)
    if len(presentation.slides) != 1 or len(presentation.slides[0].shapes) != 3:
        raise ValueError("冻结页面结构已变化")
    table_shape = shape_by_id(presentation, 3)
    picture = shape_by_id(presentation, 4)
    title = shape_by_id(presentation, 2)
    if not table_shape.has_table or len(table_shape.table.rows) != 6 or len(table_shape.table.columns) != 3:
        raise ValueError("冻结表格结构已变化")
    actual_source = tuple(tuple(cell.text for cell in row.cells) for row in table_shape.table.rows)
    if actual_source != SOURCE_ROWS:
        raise ValueError(f"冻结表格数据已变化：{actual_source!r}")

    style_before = table_style_signature(table_shape)
    picture_before = picture_signature(picture)
    title_before = (title.text, title.left, title.top, title.width, title.height)
    media_before = media_hashes(args.source)
    text_bodies = [[copy.deepcopy(cell._tc.txBody) for cell in row.cells] for row in table_shape.table.rows]

    for destination_row, source_row in enumerate(ORDER):
        for column in range(3):
            cell = table_shape.table.cell(destination_row, column)
            cell._tc.replace(cell._tc.txBody, copy.deepcopy(text_bodies[source_row][column]))

    if table_style_signature(table_shape) != style_before:
        raise ValueError("重排文本时破坏了表格结构或单元格样式")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(args.output)

    rebuilt = Presentation(args.output)
    rebuilt_table = shape_by_id(rebuilt, 3)
    actual = tuple(tuple(cell.text for cell in row.cells) for row in rebuilt_table.table.rows)
    if actual != EXPECTED_ROWS:
        raise ValueError(f"Golden 排序回读失败：{actual!r}")
    if table_style_signature(rebuilt_table) != style_before:
        raise ValueError("保存后表格样式已变化")
    if picture_signature(shape_by_id(rebuilt, 4)) != picture_before:
        raise ValueError("非目标图片的像素、几何或裁剪已变化")
    rebuilt_title = shape_by_id(rebuilt, 2)
    if (rebuilt_title.text, rebuilt_title.left, rebuilt_title.top, rebuilt_title.width, rebuilt_title.height) != title_before:
        raise ValueError("页面标题已变化")
    if media_hashes(args.output) != media_before:
        raise ValueError("媒体资产已变化")
    print(f"{args.output} {sha256(args.output)}")


if __name__ == "__main__":
    main()
