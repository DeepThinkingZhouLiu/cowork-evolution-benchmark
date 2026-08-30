#!/usr/bin/env python3
"""只为 PPTArena Case 61 构造合法 run 变体与排序/行关联/图片保持负例。"""

from __future__ import annotations

import argparse
import copy
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches


OUTPUT_NAME = "SortByScoreAndCropImage169_GroundTruth.pptx"


def output_path(root: Path, case: str) -> Path:
    target = root / case / OUTPUT_NAME
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def shape_by_id(presentation: Presentation, shape_id: int):
    matches = [shape for shape in presentation.slides[0].shapes if shape.shape_id == shape_id]
    if len(matches) != 1:
        raise ValueError(f"无法唯一定位对象 {shape_id}")
    return matches[0]


def swap_rows(table, first: int, second: int) -> None:
    first_bodies = [copy.deepcopy(table.cell(first, column)._tc.txBody) for column in range(3)]
    second_bodies = [copy.deepcopy(table.cell(second, column)._tc.txBody) for column in range(3)]
    for column in range(3):
        table.cell(first, column)._tc.replace(table.cell(first, column)._tc.txBody, second_bodies[column])
        table.cell(second, column)._tc.replace(table.cell(second, column)._tc.txBody, first_bodies[column])


def valid_alternate(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    run = shape_by_id(presentation, 3).table.cell(1, 0).text_frame.paragraphs[0].runs[0]
    if run.text != "Bob":
        raise ValueError("合法变体的 Bob run 已变化")
    duplicate = copy.deepcopy(run._r)
    run.text = "B"
    duplicate.t.text = "ob"
    run._r.addnext(duplicate)
    presentation.save(output)


def wrong_order(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    swap_rows(shape_by_id(presentation, 3).table, 2, 3)
    presentation.save(output)


def mismatched_row(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    shape_by_id(presentation, 3).table.cell(1, 1).text = "C"
    presentation.save(output)


def header_moved(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    swap_rows(shape_by_id(presentation, 3).table, 0, 1)
    presentation.save(output)


def duplicate_row(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    table = shape_by_id(presentation, 3).table
    for column in range(3):
        source = copy.deepcopy(table.cell(4, column)._tc.txBody)
        table.cell(5, column)._tc.replace(table.cell(5, column)._tc.txBody, source)
    presentation.save(output)


def wrong_fill(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    cell = shape_by_id(presentation, 3).table.cell(2, 0)
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor(255, 0, 255)
    presentation.save(output)


def changed_title(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    shape_by_id(presentation, 2).text = "Table: sorted scores"
    presentation.save(output)


def cropped_picture(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    picture = shape_by_id(presentation, 4)
    picture.crop_top = 0.125
    picture.crop_bottom = 0.125
    picture.left = 1_122_798
    picture.top = 4_027_613
    presentation.save(output)


def removed_picture(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    picture = shape_by_id(presentation, 4)
    picture._element.getparent().remove(picture._element)
    presentation.save(output)


def fake_textbox_table(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    slide = presentation.slides[0]
    table_shape = shape_by_id(presentation, 3)
    values = [[cell.text for cell in row.cells] for row in table_shape.table.rows]
    x0, y0 = table_shape.left, table_shape.top
    widths = [column.width for column in table_shape.table.columns]
    heights = [row.height for row in table_shape.table.rows]
    table_shape._element.getparent().remove(table_shape._element)
    y = y0
    for row_index, height in enumerate(heights):
        x = x0
        for column_index, width in enumerate(widths):
            box = slide.shapes.add_textbox(x, y, width, height)
            box.text = values[row_index][column_index]
            x += width
        y += height
    presentation.save(output)


def flattened_slide(golden: Path, slide_image: Path, output: Path) -> None:
    presentation = Presentation(golden)
    slide = presentation.slides[0]
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    slide.shapes.add_picture(str(slide_image), 0, 0, width=presentation.slide_width, height=presentation.slide_height)
    presentation.save(output)


def hardcoded_marker(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    box = presentation.slides[0].shapes.add_textbox(Inches(0.10), Inches(0.10), Inches(2.50), Inches(0.25))
    box.text = "solution/Golden"
    presentation.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--slide-image", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()

    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))
    valid_alternate(args.golden, output_path(args.outdir, "valid-alternate"))
    wrong_order(args.golden, output_path(args.outdir, "wrong-order"))
    mismatched_row(args.golden, output_path(args.outdir, "mismatched-row"))
    header_moved(args.golden, output_path(args.outdir, "header-moved"))
    duplicate_row(args.golden, output_path(args.outdir, "duplicate-row"))
    wrong_fill(args.golden, output_path(args.outdir, "wrong-fill"))
    changed_title(args.golden, output_path(args.outdir, "changed-title"))
    cropped_picture(args.golden, output_path(args.outdir, "cropped-picture"))
    removed_picture(args.golden, output_path(args.outdir, "removed-picture"))
    fake_textbox_table(args.golden, output_path(args.outdir, "fake-textbox-table"))
    flattened_slide(args.golden, args.slide_image, output_path(args.outdir, "flattened-slide"))
    hardcoded_marker(args.golden, output_path(args.outdir, "hardcoded-marker"))
    extra = args.outdir / "extra-output-file"
    extra.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.golden, extra / OUTPUT_NAME)
    (extra / "debug.txt").write_text("debug\n", encoding="utf-8")


if __name__ == "__main__":
    main()
