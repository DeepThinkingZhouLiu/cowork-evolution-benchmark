#!/usr/bin/env python3
"""PPTArena Case 61 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import posixpath
import re
import statistics
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from lxml import etree
from PIL import Image, ImageChops, ImageDraw, ImageStat
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "SortByScoreAndCropImage169_GroundTruth.pptx"
INPUT_NAME = "SortByScoreAndCropImage169_Test.pptx"
INPUT_SHA256 = "dd2ee0141319451ea8b33b20c06b49e2985932f302ba6de1349566a864f7995c"
TASK_ID = "cowork-evo/pptarena-061-sort-table-by-score-descending"
RUBRIC_VERSION = "pptarena-case-061-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-061-deterministic-v0.1"
SOURCE_ROWS = (
    ("Name", "Team", "Score"),
    ("Alice", "A", "82"),
    ("Bob", "B", "96"),
    ("Cara", "C", "74"),
    ("Diego", "A", "88"),
    ("Eve", "B", "91"),
)
ORDER = (0, 2, 5, 4, 1, 3)
EXPECTED_ROWS = tuple(SOURCE_ROWS[index] for index in ORDER)
SUSPICIOUS_MARKERS = ["solution/golden", "judge-result", "artifact_dir", "pptarena-case-061", "/root/"]

SPECS = {
    "R001": ("PPTX 完整且仍为原尺寸 1 页文稿", 7, "hurdle"),
    "R002": ("表头与 5 行数据按 Score 严格降序", 25, "hurdle"),
    "R003": ("Name/Team/Score 行关联完整无增删", 12, "hurdle"),
    "R004": ("原生表格结构、几何、样式和逐字符格式保持", 14, "hurdle"),
    "R005": ("标题与 Sample Signal 图片像素/位置/尺寸/裁剪保持", 16, "hurdle"),
    "R006": ("页面仍只有标题、唯一原生表格和原图片", 10, "hurdle"),
    "R007": ("整页可渲染、表格可见且表外区域保持", 10, "hurdle"),
    "R008": ("无宏、损坏关系、媒体丢失、多余输出或答案标记", 6, "positive"),
    "R009": ("无排序、行关联、可编辑性、图片或保持严重错误", -20, "penalty"),
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def safe(callback, default=""):
    try:
        return callback()
    except Exception:
        return default


def locate_source() -> Path | None:
    candidates = [Path("/app/input_files") / INPUT_NAME, Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME]
    return next((path for path in candidates if path.is_file()), None)


def media_hashes(path: Path) -> Counter[str]:
    with zipfile.ZipFile(path) as archive:
        return Counter(hashlib.sha256(archive.read(name)).hexdigest() for name in archive.namelist() if name.startswith("ppt/media/"))


def package_checks(path: Path) -> tuple[bool, dict]:
    evidence = {"broken_member": None, "missing_targets": [], "macros": []}
    try:
        with zipfile.ZipFile(path) as archive:
            evidence["broken_member"] = archive.testzip()
            names = set(archive.namelist())
            evidence["macros"] = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
            for rels_name in sorted(name for name in names if name.endswith(".rels")):
                root = etree.fromstring(archive.read(rels_name))
                if rels_name == "_rels/.rels":
                    owner = ""
                else:
                    prefix, filename = rels_name.split("/_rels/", 1)
                    owner = posixpath.join(prefix, filename[:-5])
                base = posixpath.dirname(owner)
                for relation in root:
                    if relation.attrib.get("TargetMode") == "External":
                        continue
                    target = relation.attrib.get("Target", "")
                    resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                    if resolved not in names:
                        evidence["missing_targets"].append(f"{rels_name} -> {resolved}")
        passed = evidence["broken_member"] is None and not evidence["missing_targets"] and not evidence["macros"]
        return passed, evidence
    except Exception as exc:
        evidence["exception"] = str(exc)
        return False, evidence


def font_signature(font) -> tuple:
    color = font.color
    return (
        font.name,
        round(font.size.pt, 3) if font.size else None,
        font.bold,
        font.italic,
        font.underline,
        safe(lambda: str(font.language_id)),
        safe(lambda: str(color.type)),
        safe(lambda: str(color.rgb)),
        safe(lambda: str(color.theme_color)),
        safe(lambda: round(color.brightness, 5)),
    )


def paragraph_layout(paragraph) -> tuple:
    return (
        paragraph.level,
        str(paragraph.alignment),
        str(paragraph.line_spacing),
        str(paragraph.space_before),
        str(paragraph.space_after),
    )


def character_styles(paragraph) -> tuple:
    return tuple(style for run in paragraph.runs for style in [font_signature(run.font)] * len(run.text))


def cell_style_signature(cell) -> tuple:
    return tuple((paragraph.text, character_styles(paragraph), paragraph_layout(paragraph)) for paragraph in cell.text_frame.paragraphs)


def table_style_signature(shape) -> tuple:
    return (
        shape.left,
        shape.top,
        shape.width,
        shape.height,
        tuple(row.height for row in shape.table.rows),
        tuple(column.width for column in shape.table.columns),
        etree.tostring(shape._element.graphic.graphicData.tbl.tblPr),
        tuple(etree.tostring(cell._tc.get_or_add_tcPr()) for row in shape.table.rows for cell in row.cells),
    )


def picture_signature(shape) -> tuple:
    return (
        shape.left,
        shape.top,
        shape.width,
        shape.height,
        round(shape.crop_left, 8),
        round(shape.crop_right, 8),
        round(shape.crop_top, 8),
        round(shape.crop_bottom, 8),
        hashlib.sha256(shape.image.blob).hexdigest(),
    )


def text_shape_signature(shape) -> tuple:
    return (
        shape.left,
        shape.top,
        shape.width,
        shape.height,
        tuple(cell_style_signature_proxy(paragraph) for paragraph in shape.text_frame.paragraphs),
    )


def cell_style_signature_proxy(paragraph) -> tuple:
    return paragraph.text, character_styles(paragraph), paragraph_layout(paragraph)


def shape_map(presentation: Presentation) -> dict[int, object]:
    return {shape.shape_id: shape for shape in presentation.slides[0].shapes}


def notes_signature(slide) -> tuple[str, ...]:
    try:
        return tuple(shape.text for shape in slide.notes_slide.shapes if getattr(shape, "has_text_frame", False))
    except Exception:
        return ()


def render(path: Path, outdir: Path) -> tuple[list[Path], int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "profile"
    profile.mkdir()
    completed = subprocess.run(
        ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(path)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=300,
        check=False,
    )
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return [], 0, completed.stdout[-500:]
    info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=30, check=False).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    pages = int(match.group(1)) if match else 0
    prefix = outdir / "slide"
    raster = subprocess.run(["pdftoppm", "-png", "-r", "100", str(pdf), str(prefix)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300, check=False)
    images = sorted(outdir.glob("slide-*.png")) if raster.returncode == 0 else []
    return images, pages, f"soffice={completed.returncode}, raster={raster.returncode}"


def emu_box(shape, image: Image.Image, presentation: Presentation) -> tuple[int, int, int, int]:
    return (
        max(0, int(shape.left * image.width / presentation.slide_width)),
        max(0, int(shape.top * image.height / presentation.slide_height)),
        min(image.width, int(math.ceil((shape.left + shape.width) * image.width / presentation.slide_width))),
        min(image.height, int(math.ceil((shape.top + shape.height) * image.height / presentation.slide_height))),
    )


def render_metrics(source_image: Path, candidate_image: Path, source: Presentation, candidate: Presentation, table) -> tuple[float, float]:
    with Image.open(source_image) as source_raw, Image.open(candidate_image) as candidate_raw:
        first = source_raw.convert("RGB")
        second = candidate_raw.convert("RGB")
        if first.size != second.size:
            return 0.0, 0.0
        mask = Image.new("L", first.size, 255)
        ImageDraw.Draw(mask).rectangle(emu_box(table, first, candidate), fill=0)
        difference = ImageChops.difference(first, second)
        difference.paste((0, 0, 0), mask=ImageChops.invert(mask))
        similarity = max(0.0, 1.0 - statistics.mean(ImageStat.Stat(difference).rms) / 255.0)
        crop = second.crop(emu_box(table, second, candidate)).convert("L")
        histogram = crop.histogram()
        dark_ratio = sum(histogram[:210]) / max(1, crop.width * crop.height)
        return similarity, dark_ratio


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {"criterion_id": identifier, "score": 1.0 if passed else 0.0, "evidence": evidence, "method": "deterministic", "raw": {"description": description, "weight": weight, "type": kind, "passed": bool(passed)}}


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [item["criterion_id"] for item in criteria if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0 and item["raw"]["passed"])
    penalty = sum(abs(item["raw"]["weight"]) for item in criteria if item["raw"]["weight"] < 0 and not item["raw"]["passed"])
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [sha256(path)] if path.is_file() else [], "criterion_results": criteria, "reward": reward, "passed": reward >= 0.8 and not hurdles, "hurdle_failures": hurdles}


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        return make_result(path, [result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配") for identifier in SPECS])
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        return make_result(path, [result_item(identifier, False, f"PPTX 打开失败：{exc}") for identifier in SPECS])

    criteria = []
    package_passed, package_evidence = package_checks(path)
    size_passed = len(source.slides) == len(candidate.slides) == 1 and (source.slide_width, source.slide_height) == (candidate.slide_width, candidate.slide_height)
    criteria.append(result_item("R001", package_passed and size_passed, f"页数={len(candidate.slides)}，尺寸一致={size_passed}，包={json.dumps(package_evidence, ensure_ascii=False)}"))

    tables = [shape for shape in candidate.slides[0].shapes if shape.has_table] if len(candidate.slides) == 1 else []
    table = tables[0] if len(tables) == 1 and len(tables[0].table.rows) == 6 and len(tables[0].table.columns) == 3 else None
    actual_rows = tuple(tuple(cell.text for cell in row.cells) for row in table.table.rows) if table is not None else ()
    sorted_passed = actual_rows == EXPECTED_ROWS
    criteria.append(result_item("R002", sorted_passed, f"行={actual_rows!r}"))

    data_rows = actual_rows[1:] if len(actual_rows) == 6 else ()
    scores = [int(row[2]) for row in data_rows if len(row) == 3 and row[2].isdigit()]
    integrity_passed = Counter(data_rows) == Counter(SOURCE_ROWS[1:]) and len(set(data_rows)) == 5 and scores == sorted(scores, reverse=True)
    criteria.append(result_item("R003", integrity_passed, f"多重集一致={Counter(data_rows) == Counter(SOURCE_ROWS[1:])}，唯一行={len(set(data_rows))}，分数={scores}"))

    source_shapes = shape_map(source)
    style_evidence = {}
    style_passed = False
    if table is not None:
        source_table = source_shapes[3]
        structure_ok = table_style_signature(table) == table_style_signature(source_table)
        cell_styles = []
        for destination_row, source_row in enumerate(ORDER):
            for column in range(3):
                cell_styles.append(cell_style_signature(table.table.cell(destination_row, column)) == cell_style_signature(source_table.table.cell(source_row, column)))
        style_passed = structure_ok and all(cell_styles)
        style_evidence = {"structure": structure_ok, "cell_styles": f"{sum(cell_styles)}/{len(cell_styles)}"}
    criteria.append(result_item("R004", style_passed, json.dumps(style_evidence, ensure_ascii=False)))

    candidate_shapes = shape_map(candidate) if len(candidate.slides) == 1 else {}
    title_ok = 2 in candidate_shapes and getattr(candidate_shapes[2], "has_text_frame", False) and text_shape_signature(candidate_shapes[2]) == text_shape_signature(source_shapes[2])
    picture_ok = 4 in candidate_shapes and candidate_shapes[4].shape_type == MSO_SHAPE_TYPE.PICTURE and picture_signature(candidate_shapes[4]) == picture_signature(source_shapes[4])
    notes_ok = len(candidate.slides) == 1 and notes_signature(candidate.slides[0]) == notes_signature(source.slides[0])
    media_ok = media_hashes(path) == media_hashes(source_path)
    preservation_passed = title_ok and picture_ok and notes_ok and media_ok
    criteria.append(result_item("R005", preservation_passed, f"标题={title_ok}，图片={picture_ok}，备注={notes_ok}，媒体={media_ok}"))

    full_slide_pictures = []
    if len(candidate.slides) == 1:
        page_area = candidate.slide_width * candidate.slide_height
        full_slide_pictures = [shape.name for shape in candidate.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.width * shape.height >= page_area * 0.80]
    editable_passed = table is not None and len(candidate.slides[0].shapes) == 3 and sorted(shape.shape_type for shape in candidate.slides[0].shapes) == sorted([MSO_SHAPE_TYPE.TEXT_BOX, MSO_SHAPE_TYPE.TABLE, MSO_SHAPE_TYPE.PICTURE]) and not full_slide_pictures
    criteria.append(result_item("R006", editable_passed, f"对象数={len(candidate.slides[0].shapes) if len(candidate.slides) == 1 else -1}，原生表格数={len(tables)}，整页图={full_slide_pictures}"))

    render_passed = False
    render_evidence = {"skipped": "前置条件未通过"}
    prerequisites = table is not None and all(next(item for item in criteria if item["criterion_id"] == identifier)["raw"]["passed"] for identifier in ("R001", "R002", "R003", "R004", "R005", "R006"))
    if prerequisites:
        with tempfile.TemporaryDirectory(prefix="pptarena061-judge-") as temporary:
            root = Path(temporary)
            source_images, source_pages, source_note = render(source_path, root / "source")
            candidate_images, candidate_pages, candidate_note = render(path, root / "candidate")
            similarity, dark_ratio = (0.0, 0.0)
            if len(source_images) == len(candidate_images) == 1:
                similarity, dark_ratio = render_metrics(source_images[0], candidate_images[0], source, candidate, table)
            render_passed = source_pages == candidate_pages == 1 and similarity >= 0.9995 and dark_ratio >= 0.035
            render_evidence = {"source": source_note, "candidate": candidate_note, "source_pages": source_pages, "candidate_pages": candidate_pages, "non_table_similarity": round(similarity, 6), "table_dark_ratio": round(dark_ratio, 6)}
    criteria.append(result_item("R007", render_passed, json.dumps(render_evidence, ensure_ascii=False)))

    all_text = "\n".join(shape.text for shape in candidate.slides[0].shapes if getattr(shape, "has_text_frame", False)) if len(candidate.slides) == 1 else ""
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker.lower() in all_text.lower()]
    artifact_files = sorted(str(item.relative_to(artifact_dir)) for item in artifact_dir.rglob("*") if item.is_file())
    hygiene_passed = package_passed and media_ok and artifact_files == [OUTPUT_NAME] and not markers
    criteria.append(result_item("R008", hygiene_passed, f"输出={artifact_files}，标记={markers}，媒体={media_ok}"))

    severe = not all(next(item for item in criteria if item["criterion_id"] == identifier)["raw"]["passed"] for identifier in ("R001", "R002", "R003", "R004", "R005", "R006", "R007")) or bool(markers)
    criteria.append(result_item("R009", not severe, f"严重错误={severe}"))
    order = {identifier: index for index, identifier in enumerate(SPECS)}
    criteria.sort(key=lambda item: order[item["criterion_id"]])
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    output = args.artifact_dir / OUTPUT_NAME
    if output.is_file():
        evaluation = evaluate(output, args.artifact_dir)
    else:
        evaluation = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [], "criterion_results": [], "reward": 0.0, "passed": False, "hurdle_failures": ["missing_output"]}
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
