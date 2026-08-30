#!/usr/bin/env python3
"""从 PPTArena Case 64 冻结输入构建只修改三个图表系列颜色的 Golden。"""

from __future__ import annotations

import argparse
import hashlib
import zipfile
from pathlib import Path

from pptx import Presentation


INPUT_SHA256 = "a14d9499a9bea4738c648bfb686a25e2759edc753154cd439abbef9c8444d60e"
CHART_PART = "ppt/charts/chart1.xml"
REPLACEMENTS = ((b'C83232', b'0072B2'), (b'289646', b'E69F00'), (b'F08C1E', b'009E73'))


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def member_hashes(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        return {name: hashlib.sha256(archive.read(name)).hexdigest() for name in archive.namelist()}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if not args.source.is_file() or sha256(args.source) != INPUT_SHA256:
        raise ValueError("冻结输入缺失或 SHA256 不匹配")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(args.source) as source_archive, zipfile.ZipFile(args.output, "w") as output_archive:
        names = source_archive.namelist()
        if CHART_PART not in names:
            raise ValueError("冻结图表部件缺失")
        for info in source_archive.infolist():
            data = source_archive.read(info.filename)
            if info.filename == CHART_PART:
                for old, new in REPLACEMENTS:
                    if data.count(old) != 1:
                        raise ValueError(f"原系列色 {old.decode()} 计数异常")
                    data = data.replace(old, new)
            output_archive.writestr(info, data)

    source_hashes = member_hashes(args.source)
    output_hashes = member_hashes(args.output)
    if source_hashes.keys() != output_hashes.keys():
        raise ValueError("ZIP 成员列表已变化")
    changed = [name for name in source_hashes if source_hashes[name] != output_hashes[name]]
    if changed != [CHART_PART]:
        raise ValueError(f"修改边界异常：{changed}")

    presentation = Presentation(args.output)
    charts = [shape for shape in presentation.slides[0].shapes if shape.has_chart]
    if len(presentation.slides) != 1 or len(charts) != 1 or len(charts[0].chart.series) != 3:
        raise ValueError("图表结构已变化")
    colors = tuple(str(series.format.fill.fore_color.rgb) for series in charts[0].chart.series)
    if colors != ("0072B2", "E69F00", "009E73"):
        raise ValueError(f"回读系列色异常：{colors}")
    print(f"{args.output} {sha256(args.output)}")


if __name__ == "__main__":
    main()
