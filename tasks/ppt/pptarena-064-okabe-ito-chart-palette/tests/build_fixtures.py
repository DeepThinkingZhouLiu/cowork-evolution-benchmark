#!/usr/bin/env python3
"""只为 PPTArena Case 64 构造合法重打包变体与配色/图表/保持负例。"""

from __future__ import annotations

import argparse
import copy
import shutil
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches


OUTPUT_NAME = "ColorblindPalette_GroundTruthA.pptx"
CHART_PART = "ppt/charts/chart1.xml"


def output_path(root: Path, case: str) -> Path:
    target = root / case / OUTPUT_NAME
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def rewrite_chart(source: Path, output: Path, transform) -> None:
    with zipfile.ZipFile(source) as source_archive, zipfile.ZipFile(output, "w") as output_archive:
        for info in source_archive.infolist():
            data = source_archive.read(info.filename)
            if info.filename == CHART_PART:
                data = transform(data)
            output_archive.writestr(info, data)


def replace_once(data: bytes, old: bytes, new: bytes) -> bytes:
    if data.count(old) != 1:
        raise ValueError(f"替换目标计数异常：{old!r}")
    return data.replace(old, new)


def valid_alternate(golden: Path, output: Path) -> None:
    with zipfile.ZipFile(golden) as source_archive, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_STORED) as output_archive:
        for info in source_archive.infolist():
            data = source_archive.read(info.filename)
            fresh = copy.copy(info)
            fresh.compress_type = zipfile.ZIP_STORED
            output_archive.writestr(fresh, data)


def wrong_palette_order(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    chart = next(shape.chart for shape in presentation.slides[0].shapes if shape.has_chart)
    chart.series[0].format.fill.fore_color.rgb = RGBColor(0xE6, 0x9F, 0x00)
    chart.series[1].format.fill.fore_color.rgb = RGBColor(0x00, 0x72, 0xB2)
    presentation.save(output)


def nonstandard_color(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    chart = next(shape.chart for shape in presentation.slides[0].shapes if shape.has_chart)
    chart.series[2].format.fill.fore_color.rgb = RGBColor(0x00, 0xFF, 0x00)
    presentation.save(output)


def wrong_value(golden: Path, output: Path) -> None:
    rewrite_chart(
        golden,
        output,
        lambda data: replace_once(
            data,
            b'<c:pt idx="0"><c:v>19</c:v></c:pt><c:pt idx="1"><c:v>22</c:v>',
            b'<c:pt idx="0"><c:v>20</c:v></c:pt><c:pt idx="1"><c:v>22</c:v>',
        ),
    )


def wrong_series_name(golden: Path, output: Path) -> None:
    rewrite_chart(golden, output, lambda data: replace_once(data, b'<c:v>North</c:v>', b'<c:v>Northeast</c:v>'))


def wrong_chart_type(golden: Path, output: Path) -> None:
    rewrite_chart(golden, output, lambda data: replace_once(data, b'<c:barDir val="col"/>', b'<c:barDir val="bar"/>'))


def removed_legend(golden: Path, output: Path) -> None:
    def transform(data: bytes) -> bytes:
        start = data.find(b"<c:legend>")
        end = data.find(b"</c:legend>")
        if start < 0 or end < 0:
            raise ValueError("无法定位图例")
        return data[:start] + data[end + len(b"</c:legend>"):]
    rewrite_chart(golden, output, transform)


def duplicate_chart(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    slide = presentation.slides[0]
    chart_shape = next(shape for shape in slide.shapes if shape.has_chart)
    duplicate = copy.deepcopy(chart_shape._element)
    duplicate.nvGraphicFramePr.cNvPr.id = 5
    duplicate.nvGraphicFramePr.cNvPr.name = "Duplicate Chart"
    slide.shapes._spTree.append(duplicate)
    presentation.save(output)


def overlays(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    slide = presentation.slides[0]
    chart_shape = next(shape for shape in slide.shapes if shape.has_chart)
    colors = [(0x00, 0x72, 0xB2), (0xE6, 0x9F, 0x00), (0x00, 0x9E, 0x73)]
    for index, color in enumerate(colors):
        box = slide.shapes.add_shape(1, chart_shape.left + Inches(0.5 + index), chart_shape.top + Inches(1.5), Inches(0.6), Inches(2.0))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(*color)
        box.line.fill.background()
    presentation.save(output)


def changed_instruction(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    shape = next(shape for shape in presentation.slides[0].shapes if shape.shape_id == 3)
    shape.text = "Instruction: use accessible colors."
    presentation.save(output)


def moved_chart(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    shape = next(shape for shape in presentation.slides[0].shapes if shape.has_chart)
    shape.left += Inches(0.5)
    presentation.save(output)


def flattened_slide(golden: Path, slide_image: Path, output: Path) -> None:
    presentation = Presentation(golden)
    slide = presentation.slides[0]
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    slide.shapes.add_picture(str(slide_image), 0, 0, width=presentation.slide_width, height=presentation.slide_height)
    presentation.save(output)


def hardcoded_marker(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    box = presentation.slides[0].shapes.add_textbox(Inches(0.10), Inches(0.10), Inches(2.50), Inches(0.25))
    box.text = "solution/Golden"
    presentation.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--slide-image", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()

    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))
    valid_alternate(args.golden, output_path(args.outdir, "valid-alternate"))
    wrong_palette_order(args.golden, output_path(args.outdir, "wrong-palette-order"))
    nonstandard_color(args.golden, output_path(args.outdir, "nonstandard-color"))
    wrong_value(args.golden, output_path(args.outdir, "wrong-value"))
    wrong_series_name(args.golden, output_path(args.outdir, "wrong-series-name"))
    wrong_chart_type(args.golden, output_path(args.outdir, "wrong-chart-type"))
    removed_legend(args.golden, output_path(args.outdir, "removed-legend"))
    duplicate_chart(args.golden, output_path(args.outdir, "duplicate-chart"))
    overlays(args.golden, output_path(args.outdir, "overlay-shapes"))
    changed_instruction(args.golden, output_path(args.outdir, "changed-instruction"))
    moved_chart(args.golden, output_path(args.outdir, "moved-chart"))
    flattened_slide(args.golden, args.slide_image, output_path(args.outdir, "flattened-slide"))
    hardcoded_marker(args.golden, output_path(args.outdir, "hardcoded-marker"))
    extra = args.outdir / "extra-output-file"
    extra.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.golden, extra / OUTPUT_NAME)
    (extra / "debug.txt").write_text("debug\n", encoding="utf-8")


if __name__ == "__main__":
    main()
