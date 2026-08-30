#!/usr/bin/env python3
"""PPTArena Case 64 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import posixpath
import re
import statistics
import subprocess
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image, ImageChops, ImageDraw, ImageStat
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "ColorblindPalette_GroundTruthA.pptx"
INPUT_NAME = "ColorblindPalette_TestA.pptx"
INPUT_SHA256 = "a14d9499a9bea4738c648bfb686a25e2759edc753154cd439abbef9c8444d60e"
TASK_ID = "cowork-evo/pptarena-064-okabe-ito-chart-palette"
RUBRIC_VERSION = "pptarena-case-064-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-064-deterministic-v0.1"
EXPECTED_COLORS = ("0072B2", "E69F00", "009E73")
EXPECTED_NAMES = ("North", "East", "West")
EXPECTED_CATEGORIES = ("Q1", "Q2", "Q3", "Q4")
EXPECTED_VALUES = ((19.0, 22.0, 25.0, 23.0), (15.0, 18.0, 21.0, 19.0), (12.0, 16.0, 17.0, 18.0))
SUSPICIOUS_MARKERS = ["solution/golden", "judge-result", "artifact_dir", "pptarena-case-064", "/root/"]

SPECS = {
    "R001": ("PPTX 完整且仍为原尺寸 1 页文稿", 7, "hurdle"),
    "R002": ("三个系列按顺序使用精确 Okabe–Ito 色值", 24, "hurdle"),
    "R003": ("唯一原生可编辑簇状柱形图且无覆盖", 12, "hurdle"),
    "R004": ("系列、分类、数值、标题、图例、坐标轴、标签和间距保持", 22, "hurdle"),
    "R005": ("图表几何、标题、说明文本、布局和备注保持", 14, "hurdle"),
    "R006": ("整页可渲染、新配色可见且图表外区域保持", 13, "hurdle"),
    "R007": ("无宏、损坏关系、多余输出或答案标记", 8, "positive"),
    "R008": ("无配色、数据、图表性质、可编辑性或保持严重错误", -20, "penalty"),
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def safe(callback, default=""):
    try:
        return callback()
    except Exception:
        return default


def locate_source() -> Path | None:
    candidates = [Path("/app/input_files") / INPUT_NAME, Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME]
    return next((path for path in candidates if path.is_file()), None)


def package_checks(path: Path) -> tuple[bool, dict]:
    evidence = {"broken_member": None, "missing_targets": [], "macros": []}
    try:
        with zipfile.ZipFile(path) as archive:
            evidence["broken_member"] = archive.testzip()
            names = set(archive.namelist())
            evidence["macros"] = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
            for rels_name in sorted(name for name in names if name.endswith(".rels")):
                root = etree.fromstring(archive.read(rels_name))
                if rels_name == "_rels/.rels":
                    owner = ""
                else:
                    prefix, filename = rels_name.split("/_rels/", 1)
                    owner = posixpath.join(prefix, filename[:-5])
                base = posixpath.dirname(owner)
                for relation in root:
                    if relation.attrib.get("TargetMode") == "External":
                        continue
                    target = relation.attrib.get("Target", "")
                    resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                    if resolved not in names:
                        evidence["missing_targets"].append(f"{rels_name} -> {resolved}")
        passed = evidence["broken_member"] is None and not evidence["missing_targets"] and not evidence["macros"]
        return passed, evidence
    except Exception as exc:
        evidence["exception"] = str(exc)
        return False, evidence


def font_signature(font) -> tuple:
    color = font.color
    return (
        font.name,
        round(font.size.pt, 3) if font.size else None,
        font.bold,
        font.italic,
        font.underline,
        safe(lambda: str(font.language_id)),
        safe(lambda: str(color.type)),
        safe(lambda: str(color.rgb)),
        safe(lambda: str(color.theme_color)),
    )


def text_shape_signature(shape) -> tuple:
    return (
        shape.left,
        shape.top,
        shape.width,
        shape.height,
        tuple(
            (
                paragraph.text,
                paragraph.level,
                str(paragraph.alignment),
                str(paragraph.line_spacing),
                tuple((run.text, font_signature(run.font)) for run in paragraph.runs),
            )
            for paragraph in shape.text_frame.paragraphs
        ),
    )


def chart_shapes(presentation: Presentation) -> list:
    return [shape for slide in presentation.slides for shape in slide.shapes if shape.has_chart]


def category_labels(chart) -> tuple[str, ...]:
    return tuple(str(category.label) for category in chart.plots[0].categories)


def chart_semantics(chart) -> dict:
    plot = chart.plots[0]
    data_labels = safe(lambda: plot.data_labels, None) if safe(lambda: plot.has_data_labels, False) else None
    return {
        "chart_type": str(chart.chart_type),
        "series_names": tuple(series.name for series in chart.series),
        "values": tuple(tuple(float(value) for value in series.values) for series in chart.series),
        "categories": category_labels(chart),
        "has_title": chart.has_title,
        "title": safe(lambda: chart.chart_title.text_frame.text),
        "has_legend": chart.has_legend,
        "legend_position": safe(lambda: str(chart.legend.position)),
        "legend_overlay": safe(lambda: chart.legend.include_in_layout),
        "vary_by_categories": safe(lambda: plot.vary_by_categories),
        "gap_width": safe(lambda: plot.gap_width),
        "overlap": safe(lambda: plot.overlap),
        "has_data_labels": safe(lambda: plot.has_data_labels, False),
        "data_labels": (
            safe(lambda: data_labels.show_value),
            safe(lambda: data_labels.show_category_name),
            safe(lambda: data_labels.show_series_name),
            safe(lambda: data_labels.show_legend_key),
            safe(lambda: data_labels.show_percentage),
        ) if data_labels is not None else None,
        "category_axis": (
            safe(lambda: str(chart.category_axis.axis_position)),
            safe(lambda: chart.category_axis.visible),
            safe(lambda: chart.category_axis.has_title),
            safe(lambda: chart.category_axis.tick_labels.number_format),
        ),
        "value_axis": (
            safe(lambda: str(chart.value_axis.axis_position)),
            safe(lambda: chart.value_axis.visible),
            safe(lambda: chart.value_axis.has_title),
            safe(lambda: chart.value_axis.minimum_scale, None),
            safe(lambda: chart.value_axis.maximum_scale, None),
            safe(lambda: chart.value_axis.major_unit, None),
            safe(lambda: chart.value_axis.tick_labels.number_format),
        ),
    }


def series_colors(chart) -> tuple[str, ...]:
    return tuple(safe(lambda series=series: str(series.format.fill.fore_color.rgb)) for series in chart.series)


def chart_xml_flags(path: Path) -> dict:
    with zipfile.ZipFile(path) as archive:
        chart_names = sorted(name for name in archive.namelist() if name.startswith("ppt/charts/chart") and name.endswith(".xml"))
        if len(chart_names) != 1:
            return {"chart_parts": chart_names, "dPt": -1, "ser": -1}
        root = etree.fromstring(archive.read(chart_names[0]))
        namespaces = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
        return {"chart_parts": chart_names, "dPt": len(root.xpath(".//c:dPt", namespaces=namespaces)), "ser": len(root.xpath(".//c:ser", namespaces=namespaces))}


def notes_signature(slide) -> tuple[str, ...]:
    try:
        return tuple(shape.text for shape in slide.notes_slide.shapes if getattr(shape, "has_text_frame", False))
    except Exception:
        return ()


def render(path: Path, outdir: Path) -> tuple[list[Path], int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "profile"
    profile.mkdir()
    completed = subprocess.run(
        ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(path)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=300,
        check=False,
    )
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return [], 0, completed.stdout[-500:]
    info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=30, check=False).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    pages = int(match.group(1)) if match else 0
    prefix = outdir / "slide"
    raster = subprocess.run(["pdftoppm", "-png", "-r", "100", str(pdf), str(prefix)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300, check=False)
    images = sorted(outdir.glob("slide-*.png")) if raster.returncode == 0 else []
    return images, pages, f"soffice={completed.returncode}, raster={raster.returncode}"


def emu_box(shape, image: Image.Image, presentation: Presentation) -> tuple[int, int, int, int]:
    return (
        max(0, int(shape.left * image.width / presentation.slide_width)),
        max(0, int(shape.top * image.height / presentation.slide_height)),
        min(image.width, int(math.ceil((shape.left + shape.width) * image.width / presentation.slide_width))),
        min(image.height, int(math.ceil((shape.top + shape.height) * image.height / presentation.slide_height))),
    )


def render_metrics(source_image: Path, candidate_image: Path, source: Presentation, candidate: Presentation, chart_shape) -> tuple[float, float]:
    with Image.open(source_image) as source_raw, Image.open(candidate_image) as candidate_raw:
        first = source_raw.convert("RGB")
        second = candidate_raw.convert("RGB")
        if first.size != second.size:
            return 0.0, 0.0
        mask = Image.new("L", first.size, 255)
        ImageDraw.Draw(mask).rectangle(emu_box(chart_shape, first, candidate), fill=0)
        difference = ImageChops.difference(first, second)
        difference.paste((0, 0, 0), mask=ImageChops.invert(mask))
        similarity = max(0.0, 1.0 - statistics.mean(ImageStat.Stat(difference).rms) / 255.0)
        crop = second.crop(emu_box(chart_shape, second, candidate)).convert("L")
        histogram = crop.histogram()
        dark_ratio = sum(histogram[:230]) / max(1, crop.width * crop.height)
        return similarity, dark_ratio


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {"criterion_id": identifier, "score": 1.0 if passed else 0.0, "evidence": evidence, "method": "deterministic", "raw": {"description": description, "weight": weight, "type": kind, "passed": bool(passed)}}


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [item["criterion_id"] for item in criteria if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0 and item["raw"]["passed"])
    penalty = sum(abs(item["raw"]["weight"]) for item in criteria if item["raw"]["weight"] < 0 and not item["raw"]["passed"])
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [sha256(path)] if path.is_file() else [], "criterion_results": criteria, "reward": reward, "passed": reward >= 0.8 and not hurdles, "hurdle_failures": hurdles}


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        return make_result(path, [result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配") for identifier in SPECS])
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        return make_result(path, [result_item(identifier, False, f"PPTX 打开失败：{exc}") for identifier in SPECS])

    criteria = []
    package_passed, package_evidence = package_checks(path)
    size_passed = len(source.slides) == len(candidate.slides) == 1 and (source.slide_width, source.slide_height) == (candidate.slide_width, candidate.slide_height)
    criteria.append(result_item("R001", package_passed and size_passed, f"页数={len(candidate.slides)}，尺寸一致={size_passed}，包={json.dumps(package_evidence, ensure_ascii=False)}"))

    charts = chart_shapes(candidate)
    chart_shape = charts[0] if len(charts) == 1 else None
    colors = series_colors(chart_shape.chart) if chart_shape is not None else ()
    flags = chart_xml_flags(path)
    colors_passed = colors == EXPECTED_COLORS and flags["ser"] == 3 and flags["dPt"] == 0
    criteria.append(result_item("R002", colors_passed, f"系列色={colors}，XML={flags}"))

    full_slide_pictures = []
    if len(candidate.slides) == 1:
        page_area = candidate.slide_width * candidate.slide_height
        full_slide_pictures = [shape.name for shape in candidate.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.width * shape.height >= page_area * 0.80]
    editable_passed = chart_shape is not None and len(candidate.slides[0].shapes) == 3 and not full_slide_pictures
    criteria.append(result_item("R003", editable_passed, f"原生图表数={len(charts)}，页面对象={len(candidate.slides[0].shapes) if len(candidate.slides) == 1 else -1}，整页图={full_slide_pictures}"))

    source_chart_shape = chart_shapes(source)[0]
    semantics = chart_semantics(chart_shape.chart) if chart_shape is not None else {}
    source_semantics = chart_semantics(source_chart_shape.chart)
    data_truth = semantics.get("series_names") == EXPECTED_NAMES and semantics.get("categories") == EXPECTED_CATEGORIES and semantics.get("values") == EXPECTED_VALUES
    chart_preserved = semantics == source_semantics
    criteria.append(result_item("R004", data_truth and chart_preserved, f"真值={data_truth}，图表语义保持={chart_preserved}，实际={json.dumps(semantics, ensure_ascii=False, default=str)}"))

    source_shapes = {shape.shape_id: shape for shape in source.slides[0].shapes}
    candidate_shapes = {shape.shape_id: shape for shape in candidate.slides[0].shapes} if len(candidate.slides) == 1 else {}
    geometry_ok = chart_shape is not None and (chart_shape.left, chart_shape.top, chart_shape.width, chart_shape.height) == (source_chart_shape.left, source_chart_shape.top, source_chart_shape.width, source_chart_shape.height)
    text_ok = all(shape_id in candidate_shapes and getattr(candidate_shapes[shape_id], "has_text_frame", False) and text_shape_signature(candidate_shapes[shape_id]) == text_shape_signature(source_shapes[shape_id]) for shape_id in (2, 3))
    notes_ok = len(candidate.slides) == 1 and notes_signature(candidate.slides[0]) == notes_signature(source.slides[0])
    preservation_passed = geometry_ok and text_ok and notes_ok
    criteria.append(result_item("R005", preservation_passed, f"图表几何={geometry_ok}，文本对象={text_ok}，备注={notes_ok}"))

    render_passed = False
    render_evidence = {"skipped": "前置条件未通过"}
    prerequisites = chart_shape is not None and all(next(item for item in criteria if item["criterion_id"] == identifier)["raw"]["passed"] for identifier in ("R001", "R002", "R003", "R004", "R005"))
    if prerequisites:
        with tempfile.TemporaryDirectory(prefix="pptarena064-judge-") as temporary:
            root = Path(temporary)
            source_images, source_pages, source_note = render(source_path, root / "source")
            candidate_images, candidate_pages, candidate_note = render(path, root / "candidate")
            similarity, dark_ratio = (0.0, 0.0)
            if len(source_images) == len(candidate_images) == 1:
                similarity, dark_ratio = render_metrics(source_images[0], candidate_images[0], source, candidate, chart_shape)
            render_passed = source_pages == candidate_pages == 1 and similarity >= 0.9995 and dark_ratio >= 0.04
            render_evidence = {"source": source_note, "candidate": candidate_note, "source_pages": source_pages, "candidate_pages": candidate_pages, "non_chart_similarity": round(similarity, 6), "chart_dark_ratio": round(dark_ratio, 6)}
    criteria.append(result_item("R006", render_passed, json.dumps(render_evidence, ensure_ascii=False)))

    all_text = "\n".join(shape.text for shape in candidate.slides[0].shapes if getattr(shape, "has_text_frame", False)) if len(candidate.slides) == 1 else ""
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker.lower() in all_text.lower()]
    artifact_files = sorted(str(item.relative_to(artifact_dir)) for item in artifact_dir.rglob("*") if item.is_file())
    hygiene_passed = package_passed and artifact_files == [OUTPUT_NAME] and not markers
    criteria.append(result_item("R007", hygiene_passed, f"输出={artifact_files}，标记={markers}"))

    severe = not all(next(item for item in criteria if item["criterion_id"] == identifier)["raw"]["passed"] for identifier in ("R001", "R002", "R003", "R004", "R005", "R006")) or bool(markers)
    criteria.append(result_item("R008", not severe, f"严重错误={severe}"))
    order = {identifier: index for index, identifier in enumerate(SPECS)}
    criteria.sort(key=lambda item: order[item["criterion_id"]])
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    output = args.artifact_dir / OUTPUT_NAME
    if output.is_file():
        evaluation = evaluate(output, args.artifact_dir)
    else:
        evaluation = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [], "criterion_results": [], "reward": 0.0, "passed": False, "hurdle_failures": ["missing_output"]}
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
