#!/usr/bin/env python3
"""PPTArena Case 67 专用 Golden 构建器。"""

from __future__ import annotations

import argparse
import hashlib
import shutil
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation


INPUT_SHA256 = "8e3ecec6a46f3ef3a41d764d4853359d919c840baaf5834a15e2c0814dc7e35b"
UPSTREAM_SHA256 = "b5e4479bb8b992d512e469c2b3a30ea1eef937503b901666784822400ec5309f"
TITLES = (
    "Five-Idea SmartArt Table",
    "Editable Five-Idea Table",
    "Six-Idea SmartArt Grid",
    "Five Main Points",
    "Template Design Resources",
    "Credits and Attribution",
)
TITLE_IDS = (158, 164, 199, 247, 335, 346)
READING_ORDER = {
    1: [158, 157, 117, 133, 102, 137, 138, 139, 140, 120, 132, 105, 141, 142, 143, 144, 123, 135, 108, 145, 146, 147, 148, 126, 134, 111, 149, 150, 151, 152, 129, 136, 114, 153, 154, 155, 156],
    2: [164, 163, 162],
    3: [199, 186, 168, 188, 187, 171, 190, 189, 174, 192, 191, 177, 194, 193, 180, 196, 195, 183, 198, 197],
    4: [247, 246, 203, 234, 237, 236, 231, 239, 238, 232, 241, 240, 235, 243, 242, 233, 245, 244],
    5: [335, 275, 273, 251, 254, 257, 260, 276, 277, 269, 282, 271, 270, 263, 278, 266, 279, 322, 325, 328, 272, 331, 332, 333, 334, 274],
    6: [346, 343, 339, 345, 341, 344, 342, 340],
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def shape_map(slide) -> dict:
    return {shape.shape_id: shape for shape in slide.shapes}


def apply_titles_and_order(presentation: Presentation) -> None:
    for slide_number, (slide, title, title_id) in enumerate(zip(presentation.slides, TITLES, TITLE_IDS, strict=True), 1):
        mapping = shape_map(slide)
        if set(mapping) != set(READING_ORDER[slide_number]) or title_id not in mapping:
            raise ValueError(f"第 {slide_number} 页对象 ID 与已审计顺序不匹配")
        title_shape = mapping[title_id]
        title_shape.text = title
        title_shape.left = -9144000
        title_shape.top = -9144000
        title_shape.width = 9144
        title_shape.height = 9144
        properties = title_shape._element.nvSpPr.cNvPr
        properties.set("name", f"Accessibility Title {slide_number}: {title}")
        properties.set("title", title)
        properties.set("descr", f"Unique machine-readable title for slide {slide_number}: {title}")
        for shape_id in READING_ORDER[slide_number]:
            slide.shapes._spTree.append(mapping[shape_id]._element)


def strip_direct_typefaces(path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="pptarena067-fonts-") as temporary:
        temporary_path = Path(temporary) / path.name
        with zipfile.ZipFile(path) as reader, zipfile.ZipFile(temporary_path, "w", compression=zipfile.ZIP_DEFLATED) as writer:
            for info in reader.infolist():
                data = reader.read(info.filename)
                if info.filename.endswith(".xml") and info.filename.startswith(("ppt/slides/", "ppt/notesSlides/", "ppt/slideMasters/", "ppt/slideLayouts/")):
                    root = etree.fromstring(data)
                    for node in list(root.xpath(".//*[local-name()='rPr' or local-name()='defRPr' or local-name()='endParaRPr']/*[local-name()='latin' or local-name()='ea' or local-name()='cs' or local-name()='sym']")):
                        node.getparent().remove(node)
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                writer.writestr(info, data)
        shutil.copy2(temporary_path, path)


def direct_typeface_count(path: Path) -> int:
    count = 0
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.endswith(".xml") and name.startswith(("ppt/slides/", "ppt/notesSlides/", "ppt/slideMasters/", "ppt/slideLayouts/")):
                root = etree.fromstring(archive.read(name))
                count += len(root.xpath(".//*[local-name()='rPr' or local-name()='defRPr' or local-name()='endParaRPr']/*[local-name()='latin' or local-name()='ea' or local-name()='cs' or local-name()='sym']"))
    return count


def build(source: Path, upstream: Path, output: Path) -> None:
    if sha256(source) != INPUT_SHA256 or sha256(upstream) != UPSTREAM_SHA256:
        raise ValueError("冻结输入或上游 Ground Truth SHA256 不匹配")
    presentation = Presentation(upstream)
    if len(presentation.slides) != 6:
        raise ValueError("上游 Ground Truth 不再是 6 页")
    apply_titles_and_order(presentation)
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    strip_direct_typefaces(output)
    reopened = Presentation(output)
    for slide_number, slide in enumerate(reopened.slides, 1):
        if [shape.shape_id for shape in slide.shapes] != READING_ORDER[slide_number]:
            raise ValueError(f"第 {slide_number} 页阅读顺序保存后不匹配")
        if slide.shapes[0].text != TITLES[slide_number - 1]:
            raise ValueError(f"第 {slide_number} 页唯一标题不匹配")
    if direct_typeface_count(output) != 0:
        raise ValueError("产物仍含直接字体族覆盖")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--upstream", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.source, args.upstream, args.output)
    print(f"{args.output}\nsha256={sha256(args.output)}")


if __name__ == "__main__":
    main()
