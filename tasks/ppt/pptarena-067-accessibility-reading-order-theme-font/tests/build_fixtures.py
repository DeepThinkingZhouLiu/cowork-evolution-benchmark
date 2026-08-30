#!/usr/bin/env python3
"""PPTArena Case 67 专用 Judge 校准样例。"""

from __future__ import annotations

import argparse
import shutil
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt


OUTPUT_NAME = "AccessibilityMasterCleanup_GroundTruthA.pptx"


def output_path(root: Path, case: str) -> Path:
    path = root / case / OUTPUT_NAME
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def save_mutation(golden: Path, output: Path, callback) -> None:
    presentation = Presentation(golden)
    callback(presentation)
    presentation.save(output)


def repackage(golden: Path, output: Path) -> None:
    with zipfile.ZipFile(golden) as reader, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_STORED) as writer:
        for info in reader.infolist():
            writer.writestr(info.filename, reader.read(info.filename))


def add_direct_font(golden: Path, output: Path) -> None:
    with zipfile.ZipFile(golden) as reader, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as writer:
        changed = False
        for info in reader.infolist():
            data = reader.read(info.filename)
            if info.filename == "ppt/slides/slide2.xml":
                root = etree.fromstring(data)
                run_properties = root.xpath(".//*[local-name()='t' and text()='Idea 1']/preceding-sibling::*[local-name()='rPr'][1]")
                if len(run_properties) != 1:
                    raise ValueError("未定位第 2 页 Idea 1 表头")
                node = etree.SubElement(run_properties[0], "{http://schemas.openxmlformats.org/drawingml/2006/main}latin")
                node.set("typeface", "DM Sans")
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                changed = True
            writer.writestr(info, data)
        if not changed:
            raise ValueError("字体变异没有发生")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--upstream", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()
    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))
    shutil.copy2(args.upstream, output_path(args.outdir, "upstream-uncorrected"))
    repackage(args.golden, output_path(args.outdir, "valid-repackage"))

    def missing_title(prs):
        shape = prs.slides[0].shapes[0]
        shape._element.getparent().remove(shape._element)
    save_mutation(args.golden, output_path(args.outdir, "missing-title"), missing_title)

    def duplicate_title(prs):
        prs.slides[1].shapes[0].text = prs.slides[0].shapes[0].text
    save_mutation(args.golden, output_path(args.outdir, "duplicate-title"), duplicate_title)

    def visible_title(prs):
        title = prs.slides[2].shapes[0]
        title.left = Inches(1)
        title.top = Inches(1)
        title.width = Inches(4)
        title.height = Inches(0.5)
    save_mutation(args.golden, output_path(args.outdir, "visible-title"), visible_title)

    def wrong_order(prs):
        slide = prs.slides[2]
        first_content = slide.shapes[1]
        last_content = slide.shapes[-1]
        slide.shapes._spTree.insert(3, last_content._element)
        slide.shapes._spTree.append(first_content._element)
    save_mutation(args.golden, output_path(args.outdir, "wrong-order"), wrong_order)

    add_direct_font(args.golden, output_path(args.outdir, "direct-font-override"))

    def altered_size(prs):
        for run in prs.slides[4].shapes[1].text_frame.paragraphs[0].runs:
            run.font.size = Pt(38)
    save_mutation(args.golden, output_path(args.outdir, "altered-font-size"), altered_size)

    def moved_content(prs):
        prs.slides[3].shapes[2].left += Inches(0.2)
    save_mutation(args.golden, output_path(args.outdir, "moved-content"), moved_content)

    def changed_text(prs):
        prs.slides[5].shapes[-1].text = "Design complete."
    save_mutation(args.golden, output_path(args.outdir, "changed-text"), changed_text)

    def flattened(prs):
        slide = prs.slides[0]
        for shape in list(slide.shapes):
            shape._element.getparent().remove(shape._element)
        image = args.outdir / "flatten-placeholder.png"
        from PIL import Image
        Image.new("RGB", (1600, 900), "white").save(image)
        slide.shapes.add_picture(str(image), 0, 0, width=prs.slide_width, height=prs.slide_height)
    save_mutation(args.golden, output_path(args.outdir, "flattened"), flattened)

    def hardcoded(prs):
        box = prs.slides[0].shapes.add_textbox(Inches(0.1), Inches(0.1), Inches(2.0), Inches(0.2))
        box.text = "solution/Golden"
    save_mutation(args.golden, output_path(args.outdir, "hardcoded-marker"), hardcoded)

    extra = args.outdir / "extra-output"
    extra.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.golden, extra / OUTPUT_NAME)
    (extra / "debug.txt").write_text("debug\n", encoding="utf-8")


if __name__ == "__main__":
    main()
