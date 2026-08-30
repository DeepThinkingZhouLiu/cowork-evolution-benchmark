#!/usr/bin/env python3
"""PPTArena Case 67 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import posixpath
import re
import statistics
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from lxml import etree
from PIL import Image, ImageChops, ImageDraw, ImageStat
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "AccessibilityMasterCleanup_GroundTruthA.pptx"
INPUT_NAME = "AccessibilityMasterCleanup_TestA.pptx"
INPUT_SHA256 = "8e3ecec6a46f3ef3a41d764d4853359d919c840baaf5834a15e2c0814dc7e35b"
TASK_ID = "cowork-evo/pptarena-067-accessibility-reading-order-theme-font"
RUBRIC_VERSION = "pptarena-case-067-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-067-deterministic-v0.1"
TITLES = (
    "Five-Idea SmartArt Table",
    "Editable Five-Idea Table",
    "Six-Idea SmartArt Grid",
    "Five Main Points",
    "Template Design Resources",
    "Credits and Attribution",
)
READING_ORDER = {
    1: [157, 117, 133, 102, 137, 138, 139, 140, 120, 132, 105, 141, 142, 143, 144, 123, 135, 108, 145, 146, 147, 148, 126, 134, 111, 149, 150, 151, 152, 129, 136, 114, 153, 154, 155, 156],
    2: [163, 162],
    3: [186, 168, 188, 187, 171, 190, 189, 174, 192, 191, 177, 194, 193, 180, 196, 195, 183, 198, 197],
    4: [246, 203, 234, 237, 236, 231, 239, 238, 232, 241, 240, 235, 243, 242, 233, 245, 244],
    5: [275, 273, 251, 254, 257, 260, 276, 277, 269, 282, 271, 270, 263, 278, 266, 279, 322, 325, 328, 272, 331, 332, 333, 334, 274],
    6: [343, 339, 345, 341, 344, 342, 340],
}
SPECS = {
    "R001": ("PPTX 完整且严格保留尺寸和 6 页结构", 7, "hurdle"),
    "R002": ("6 个唯一、不可见且带元数据的机器标题", 18, "hurdle"),
    "R003": ("6 页均使用任务专属的逻辑阅读顺序", 20, "hurdle"),
    "R004": ("主题 Arial 且零直接字体族覆盖", 18, "hurdle"),
    "R005": ("原对象和原生可编辑结构完整保留", 15, "hurdle"),
    "R006": ("文字格式、备注和媒体资产保持", 12, "hurdle"),
    "R007": ("6 页可渲染且非文字区域保持", 6, "hurdle"),
    "R008": ("无宏、坏关系、多余输出或答案标记", 4, "positive"),
    "R009": ("无标题、顺序、字体、内容、几何或压平严重错误", -20, "penalty"),
}
SUSPICIOUS_MARKERS = ["solution/golden", "judge-result", "artifact_dir", "pptarena-case-067", "/root/"]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def safe(callback, default=""):
    try:
        return callback()
    except Exception:
        return default


def locate_source() -> Path | None:
    candidates = [Path("/app/input_files") / INPUT_NAME, Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME]
    return next((path for path in candidates if path.is_file()), None)


def canonical(element) -> bytes:
    return etree.tostring(element, method="c14n") if element is not None else b""


def package_checks(path: Path) -> tuple[bool, dict]:
    evidence = {"broken_member": None, "missing_targets": [], "macros": []}
    try:
        with zipfile.ZipFile(path) as archive:
            evidence["broken_member"] = archive.testzip()
            names = set(archive.namelist())
            evidence["macros"] = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
            for rels_name in sorted(name for name in names if name.endswith(".rels")):
                root = etree.fromstring(archive.read(rels_name))
                if rels_name == "_rels/.rels":
                    owner = ""
                else:
                    prefix, filename = rels_name.split("/_rels/", 1)
                    owner = posixpath.join(prefix, filename[:-5])
                base = posixpath.dirname(owner)
                for relation in root:
                    if relation.attrib.get("TargetMode") == "External":
                        continue
                    target = relation.attrib.get("Target", "")
                    resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                    if resolved not in names:
                        evidence["missing_targets"].append(f"{rels_name} -> {resolved}")
        return evidence["broken_member"] is None and not evidence["missing_targets"] and not evidence["macros"], evidence
    except Exception as exc:
        evidence["exception"] = str(exc)
        return False, evidence


def direct_typeface_evidence(path: Path) -> tuple[int, dict[str, int]]:
    total = 0
    by_part: dict[str, int] = {}
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.endswith(".xml") or not name.startswith(("ppt/slides/", "ppt/notesSlides/", "ppt/slideMasters/", "ppt/slideLayouts/")):
                continue
            root = etree.fromstring(archive.read(name))
            count = len(root.xpath(".//*[local-name()='rPr' or local-name()='defRPr' or local-name()='endParaRPr']/*[local-name()='latin' or local-name()='ea' or local-name()='cs' or local-name()='sym']"))
            if count:
                by_part[name] = count
                total += count
    return total, by_part


def theme_major_latin(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        root = etree.fromstring(archive.read("ppt/theme/theme1.xml"))
        values = root.xpath(".//*[local-name()='fontScheme']/*[local-name()='majorFont']/*[local-name()='latin']/@typeface")
    return values[0] if len(values) == 1 else ""


def title_metadata(shape) -> tuple[str, str, str]:
    nodes = shape._element.xpath(".//*[local-name()='cNvPr']")
    if len(nodes) != 1:
        return "", "", ""
    node = nodes[0]
    return node.attrib.get("name", ""), node.attrib.get("title", ""), node.attrib.get("descr", "")


def title_and_order_checks(source: Presentation, candidate: Presentation) -> tuple[bool, bool, list[dict]]:
    title_ok = True
    order_ok = True
    details = []
    observed_titles = []
    for slide_number, (source_slide, candidate_slide, expected_title) in enumerate(zip(source.slides, candidate.slides, TITLES, strict=False), 1):
        source_ids = {shape.shape_id for shape in source_slide.shapes}
        candidate_ids = [shape.shape_id for shape in candidate_slide.shapes]
        extras = [shape for shape in candidate_slide.shapes if shape.shape_id not in source_ids]
        current_title_ok = len(extras) == 1 and bool(candidate_ids) and extras[0].shape_id == candidate_ids[0]
        title_shape = extras[0] if len(extras) == 1 else None
        if title_shape is not None:
            observed_titles.append(safe(lambda: title_shape.text))
            name, metadata_title, description = title_metadata(title_shape)
            offscreen = title_shape.left + title_shape.width <= 0 or title_shape.top + title_shape.height <= 0 or title_shape.left >= candidate.slide_width or title_shape.top >= candidate.slide_height
            current_title_ok = current_title_ok and safe(lambda: title_shape.text) == expected_title and offscreen
            current_title_ok = current_title_ok and expected_title in name and metadata_title == expected_title and expected_title in description
        expected_order = ([title_shape.shape_id] if title_shape is not None else []) + READING_ORDER[slide_number]
        current_order_ok = candidate_ids == expected_order
        title_ok = title_ok and current_title_ok
        order_ok = order_ok and current_order_ok
        details.append({"slide": slide_number, "title": safe(lambda: title_shape.text) if title_shape is not None else None, "extra_ids": [shape.shape_id for shape in extras], "offscreen": bool(title_shape is not None and offscreen), "title_ok": current_title_ok, "order_ok": current_order_ok, "order": candidate_ids})
    title_ok = title_ok and len(observed_titles) == 6 and len(set(observed_titles)) == 6 and tuple(observed_titles) == TITLES
    return title_ok, order_ok, details


def font_signature(font) -> tuple:
    color = font.color
    return (
        round(font.size.pt, 4) if font.size else None,
        font.bold,
        font.italic,
        font.underline,
        safe(lambda: str(font.language_id)),
        safe(lambda: str(color.type)),
        safe(lambda: str(color.rgb)),
        safe(lambda: str(color.theme_color)),
        safe(lambda: round(color.brightness, 5)),
    )


def character_styles(paragraph) -> tuple:
    return tuple(style for run in paragraph.runs for style in [font_signature(run.font)] * len(run.text))


def text_frame_signature(frame) -> tuple:
    return (
        frame.margin_left,
        frame.margin_right,
        frame.margin_top,
        frame.margin_bottom,
        str(frame.vertical_anchor),
        frame.word_wrap,
        str(frame.auto_size),
        tuple((paragraph.text, character_styles(paragraph), canonical(paragraph._p.pPr)) for paragraph in frame.paragraphs),
    )


def text_units(shape, prefix="") -> tuple:
    current = f"{prefix}/{shape.shape_id}"
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return tuple(unit for child in shape.shapes for unit in text_units(child, current))
    if getattr(shape, "has_table", False):
        units = []
        for row_index, row in enumerate(shape.table.rows):
            for column_index, cell in enumerate(row.cells):
                units.append((f"{current}/cell/{row_index}/{column_index}", text_frame_signature(cell.text_frame), canonical(cell._tc.tcPr)))
        return tuple(units)
    if getattr(shape, "has_text_frame", False):
        return ((current, text_frame_signature(shape.text_frame), b""),)
    return ()


def notes_signature(slide) -> tuple[str, ...]:
    try:
        return tuple(shape.text for shape in slide.notes_slide.shapes if getattr(shape, "has_text_frame", False))
    except Exception:
        return ()


def media_hashes(path: Path) -> Counter:
    with zipfile.ZipFile(path) as archive:
        return Counter(hashlib.sha256(archive.read(name)).hexdigest() for name in archive.namelist() if name.startswith("ppt/media/") and not name.endswith("/"))


def shape_structure_and_text(source: Presentation, candidate: Presentation) -> tuple[bool, bool, dict]:
    structure_ok = True
    text_style_ok = True
    detail = {"slides": [], "full_slide_pictures": []}
    for slide_number, (source_slide, candidate_slide) in enumerate(zip(source.slides, candidate.slides, strict=False), 1):
        source_map = {shape.shape_id: shape for shape in source_slide.shapes}
        candidate_map = {shape.shape_id: shape for shape in candidate_slide.shapes}
        missing = sorted(set(source_map) - set(candidate_map))
        extras = sorted(set(candidate_map) - set(source_map))
        slide_structure = not missing and len(extras) == 1
        for shape_id, source_shape in source_map.items():
            candidate_shape = candidate_map.get(shape_id)
            if candidate_shape is None:
                continue
            source_geometry = (str(source_shape.shape_type), source_shape.left, source_shape.top, source_shape.width, source_shape.height, safe(lambda: source_shape.rotation))
            candidate_geometry = (str(candidate_shape.shape_type), candidate_shape.left, candidate_shape.top, candidate_shape.width, candidate_shape.height, safe(lambda: candidate_shape.rotation))
            slide_structure = slide_structure and source_geometry == candidate_geometry
            text_style_ok = text_style_ok and text_units(source_shape) == text_units(candidate_shape)
        for shape in candidate_slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.width * shape.height >= candidate.slide_width * candidate.slide_height * 0.8:
                detail["full_slide_pictures"].append((slide_number, shape.shape_id))
        structure_ok = structure_ok and slide_structure
        detail["slides"].append({"slide": slide_number, "missing": missing, "extras": extras, "structure": slide_structure})
    structure_ok = structure_ok and not detail["full_slide_pictures"]
    return structure_ok, text_style_ok, detail


def render(path: Path, outdir: Path) -> tuple[list[Path], int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "profile"
    profile.mkdir()
    completed = subprocess.run(
        ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(path)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=300,
        check=False,
    )
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return [], 0, completed.stdout[-500:]
    info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=30, check=False).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    pages = int(match.group(1)) if match else 0
    prefix = outdir / "slide"
    raster = subprocess.run(["pdftoppm", "-png", "-r", "100", str(pdf), str(prefix)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300, check=False)
    images = sorted(outdir.glob("slide-*.png")) if raster.returncode == 0 else []
    return images, pages, f"soffice={completed.returncode}, raster={raster.returncode}"


def mask_shape(draw: ImageDraw.ImageDraw, shape, presentation: Presentation, image: Image.Image) -> None:
    contains_text = getattr(shape, "has_text_frame", False) or getattr(shape, "has_table", False)
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        contains_text = contains_text or any(getattr(child, "has_text_frame", False) or getattr(child, "has_table", False) for child in shape.shapes)
    if not contains_text:
        return
    left = max(0, int(shape.left * image.width / presentation.slide_width) - 4)
    top = max(0, int(shape.top * image.height / presentation.slide_height) - 4)
    right = min(image.width, int(math.ceil((shape.left + shape.width) * image.width / presentation.slide_width)) + 4)
    bottom = min(image.height, int(math.ceil((shape.top + shape.height) * image.height / presentation.slide_height)) + 4)
    if right > left and bottom > top:
        draw.rectangle((left, top, right, bottom), fill=0)


def masked_render_similarity(source_images: list[Path], candidate_images: list[Path], source: Presentation, candidate: Presentation) -> float:
    scores = []
    for index, (source_path, candidate_path) in enumerate(zip(source_images, candidate_images, strict=True)):
        with Image.open(source_path) as source_raw, Image.open(candidate_path) as candidate_raw:
            first = source_raw.convert("RGB")
            second = candidate_raw.convert("RGB")
            if first.size != second.size:
                scores.append(0.0)
                continue
            mask = Image.new("L", first.size, 255)
            draw = ImageDraw.Draw(mask)
            for shape in source.slides[index].shapes:
                mask_shape(draw, shape, source, first)
            for shape in candidate.slides[index].shapes:
                mask_shape(draw, shape, candidate, second)
            difference = ImageChops.difference(first, second)
            difference.paste((0, 0, 0), mask=ImageChops.invert(mask))
            scores.append(max(0.0, 1.0 - statistics.mean(ImageStat.Stat(difference).rms) / 255.0))
    return min(scores, default=0.0)


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {"criterion_id": identifier, "score": 1.0 if passed else 0.0, "evidence": evidence, "method": "deterministic", "raw": {"description": description, "weight": weight, "type": kind, "passed": bool(passed)}}


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [item["criterion_id"] for item in criteria if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0 and item["raw"]["passed"])
    penalty = sum(abs(item["raw"]["weight"]) for item in criteria if item["raw"]["weight"] < 0 and not item["raw"]["passed"])
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [sha256(path)] if path.is_file() else [], "criterion_results": criteria, "reward": reward, "passed": reward >= 0.8 and not hurdles, "hurdle_failures": hurdles}


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        return make_result(path, [result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配") for identifier in SPECS])
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        return make_result(path, [result_item(identifier, False, f"PPTX 打开失败：{exc}") for identifier in SPECS])

    criteria = []
    package_ok, package_evidence = package_checks(path)
    size_ok = len(source.slides) == len(candidate.slides) == 6 and (source.slide_width, source.slide_height) == (candidate.slide_width, candidate.slide_height)
    criteria.append(result_item("R001", package_ok and size_ok, f"页数={len(candidate.slides)}，尺寸一致={size_ok}，包={json.dumps(package_evidence, ensure_ascii=False)}"))

    title_ok, order_ok, accessibility_details = title_and_order_checks(source, candidate)
    criteria.append(result_item("R002", title_ok, json.dumps(accessibility_details, ensure_ascii=False)))
    criteria.append(result_item("R003", order_ok, json.dumps([{"slide": item["slide"], "order_ok": item["order_ok"]} for item in accessibility_details], ensure_ascii=False)))

    try:
        direct_count, direct_parts = direct_typeface_evidence(path)
        theme_font = theme_major_latin(path)
        font_ok = direct_count == 0 and theme_font.casefold() == "arial"
        font_evidence = f"主题主字体={theme_font!r}，直接覆盖={direct_count}，分布={json.dumps(direct_parts, ensure_ascii=False)}"
    except Exception as exc:
        font_ok = False
        font_evidence = f"字体审计失败：{exc}"
    criteria.append(result_item("R004", font_ok, font_evidence))

    structure_ok, text_style_ok, structure_detail = shape_structure_and_text(source, candidate)
    media_ok = safe(lambda: media_hashes(source_path) == media_hashes(path), False)
    notes_ok = len(source.slides) == len(candidate.slides) and all(notes_signature(a) == notes_signature(b) for a, b in zip(source.slides, candidate.slides, strict=False))
    criteria.append(result_item("R005", structure_ok, json.dumps(structure_detail, ensure_ascii=False)))
    criteria.append(result_item("R006", text_style_ok and media_ok and notes_ok, f"文字及非字体样式={text_style_ok}，媒体={media_ok}，备注={notes_ok}"))

    render_ok = False
    render_evidence = "未渲染"
    try:
        with tempfile.TemporaryDirectory(prefix="pptarena067-render-") as temporary:
            root = Path(temporary)
            source_images, source_pages, source_log = render(source_path, root / "source")
            candidate_images, candidate_pages, candidate_log = render(path, root / "candidate")
            if source_pages == candidate_pages == 6 and len(source_images) == len(candidate_images) == 6:
                similarity = masked_render_similarity(source_images, candidate_images, source, candidate)
                # 本题要求全文本从 DM Sans 回退到主题 Arial，因此文字边界附近
                # 必然存在可见变化。对所有文字对象遮罩后，非文字区域阈值按人工 Golden
                # 的实测下界 0.958 校准；对象移动和压平另由 R005 硬门槛捕获。
                render_ok = similarity >= 0.95
                render_evidence = f"非文字区域最低相似度={similarity:.6f}"
            else:
                render_evidence = f"页数={source_pages}/{candidate_pages}，{source_log}，{candidate_log}"
    except Exception as exc:
        render_evidence = f"渲染异常：{exc}"
    criteria.append(result_item("R007", render_ok, render_evidence))

    delivered = sorted(str(item.relative_to(artifact_dir)) for item in artifact_dir.rglob("*") if item.is_file())
    combined_text = "\n".join(shape.text for slide in candidate.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False)).lower()
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker in combined_text]
    hygiene_ok = delivered == [OUTPUT_NAME] and package_ok and not markers
    criteria.append(result_item("R008", hygiene_ok, f"文件={delivered}，标记={markers}"))

    severe = not (title_ok and order_ok and font_ok and structure_ok and text_style_ok and media_ok and notes_ok)
    criteria.append(result_item("R009", not severe, f"严重错误={severe}"))
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    path = args.artifact_dir / OUTPUT_NAME
    result = evaluate(path, args.artifact_dir) if path.is_file() else make_result(path, [result_item(identifier, False, f"缺失指定交付物：{OUTPUT_NAME}") for identifier in SPECS])
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{result['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
