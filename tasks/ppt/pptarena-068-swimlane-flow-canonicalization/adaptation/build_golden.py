#!/usr/bin/env python3
"""PPTArena Case 68 专用 Golden 修正器。"""

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

from lxml import etree
from pptx import Presentation


INPUT_SHA256 = "c5f0e852c29a8347cb29de3897d2c16cb11c84c5a9cd32dfec0e69444c756982"
UPSTREAM_SHA256 = "1e12abfd956b0be22f9e8d1bb73cb8018e7aec992a8010afe67a504cb5103b68"
CONNECTOR_IDS = (21, 22, 23, 24, 25)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def shape_map(presentation: Presentation) -> dict:
    return {shape.shape_id: shape for shape in presentation.slides[0].shapes}


def replace_exact_text(shape, expected: str, replacement: str) -> None:
    if shape.text != expected:
        raise ValueError(f"变异目标文字不匹配：{shape.text!r} != {expected!r}")
    runs = [run for paragraph in shape.text_frame.paragraphs for run in paragraph.runs]
    if not runs or sum(len(run.text) for run in runs) != len(replacement):
        raise ValueError("变异目标 run 长度与替换文字不兼容")
    offset = 0
    for run in runs:
        length = len(run.text)
        run.text = replacement[offset : offset + length]
        offset += length


def add_terminal_arrow(shape) -> None:
    line_nodes = shape._element.xpath("./*[local-name()='spPr']/*[local-name()='ln']")
    if len(line_nodes) != 1:
        raise ValueError(f"连接符 {shape.shape_id} 缺少唯一线样式")
    line = line_nodes[0]
    for node in list(line.xpath("./*[local-name()='tailEnd']")):
        line.remove(node)
    tail = etree.SubElement(line, "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")


def build(source: Path, upstream: Path, output: Path) -> None:
    if sha256(source) != INPUT_SHA256 or sha256(upstream) != UPSTREAM_SHA256:
        raise ValueError("冻结输入或上游 Ground Truth SHA256 不匹配")
    presentation = Presentation(upstream)
    if len(presentation.slides) != 1:
        raise ValueError("上游 Ground Truth 不再是单页")
    mapping = shape_map(presentation)
    replace_exact_text(mapping[12], "Ok?", "OK?")
    # 大写 K 比小写 k 稍宽；禁止这个三字符判定标签在菱形内换行。
    mapping[12].text_frame.word_wrap = False
    replace_exact_text(mapping[16], "2", "4")
    replace_exact_text(mapping[18], "4", "2")
    for shape_id in CONNECTOR_IDS:
        add_terminal_arrow(mapping[shape_id])
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    reopened = Presentation(output)
    mapping = shape_map(reopened)
    if mapping[12].text != "OK?" or mapping[12].text_frame.word_wrap is not False or mapping[16].text != "4" or mapping[18].text != "2":
        raise ValueError("大小写或徽章修正保存后失效")
    for shape_id in CONNECTOR_IDS:
        values = mapping[shape_id]._element.xpath("./*[local-name()='spPr']/*[local-name()='ln']/*[local-name()='tailEnd']/@type")
        if values != ["triangle"]:
            raise ValueError(f"连接符 {shape_id} 末端箭头保存后失效")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--upstream", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.source, args.upstream, args.output)
    print(f"{args.output}\nsha256={sha256(args.output)}")


if __name__ == "__main__":
    main()
