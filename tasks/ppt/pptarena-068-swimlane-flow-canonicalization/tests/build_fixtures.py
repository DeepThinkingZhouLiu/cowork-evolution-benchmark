#!/usr/bin/env python3
"""PPTArena Case 68 专用 Judge 校准样例。"""

from __future__ import annotations

import argparse
import shutil
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches


OUTPUT_NAME = "SwimlaneFlowCanonicalization_GroundTruthA.pptx"


def output_path(root: Path, case: str) -> Path:
    path = root / case / OUTPUT_NAME
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def save_mutation(golden: Path, output: Path, callback) -> None:
    presentation = Presentation(golden)
    callback(presentation)
    presentation.save(output)


def by_text(presentation: Presentation, text: str):
    matches = [shape for shape in presentation.slides[0].shapes if getattr(shape, "has_text_frame", False) and shape.text.strip() == text]
    if len(matches) != 1:
        raise ValueError(f"文字对象不唯一：{text!r}")
    return matches[0]


def connectors(presentation: Presentation) -> list:
    return [shape for shape in presentation.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.LINE]


def repackage(golden: Path, output: Path) -> None:
    with zipfile.ZipFile(golden) as reader, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_STORED) as writer:
        for info in reader.infolist():
            writer.writestr(info.filename, reader.read(info.filename))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--upstream", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()
    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))
    shutil.copy2(args.upstream, output_path(args.outdir, "upstream-uncorrected"))
    repackage(args.golden, output_path(args.outdir, "valid-repackage"))

    save_mutation(args.golden, output_path(args.outdir, "wrong-lane"), lambda prs: setattr(by_text(prs, "Step 4: Reject"), "top", Inches(1.6)))
    save_mutation(args.golden, output_path(args.outdir, "unequal-box"), lambda prs: setattr(by_text(prs, "Step 2: Validate"), "width", Inches(2.9)))
    save_mutation(args.golden, output_path(args.outdir, "badge-mismatch"), lambda prs: setattr(by_text(prs, "2").text_frame.paragraphs[0].runs[0], "text", "4"))

    def straight_connector(prs):
        geometry = connectors(prs)[0]._element.xpath("./*[local-name()='spPr']/*[local-name()='prstGeom']")[0]
        geometry.set("prst", "line")
    save_mutation(args.golden, output_path(args.outdir, "straight-connector"), straight_connector)

    def missing_arrow(prs):
        line = connectors(prs)[0]._element.xpath("./*[local-name()='spPr']/*[local-name()='ln']")[0]
        for node in list(line.xpath("./*[local-name()='tailEnd']")):
            line.remove(node)
    save_mutation(args.golden, output_path(args.outdir, "missing-arrow"), missing_arrow)

    def wrong_flow(prs):
        end = connectors(prs)[0]._element.xpath(".//*[local-name()='endCxn']")[0]
        end.set("id", str(by_text(prs, "Step 3: Transform").shape_id))
    save_mutation(args.golden, output_path(args.outdir, "wrong-flow"), wrong_flow)

    def changed_fill(prs):
        by_text(prs, "Step 3: Transform").fill.fore_color.rgb = RGBColor(255, 0, 0)
    save_mutation(args.golden, output_path(args.outdir, "changed-fill"), changed_fill)

    def image_removed(prs):
        picture = next(shape for shape in prs.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
        picture._element.getparent().remove(picture._element)
    save_mutation(args.golden, output_path(args.outdir, "image-removed"), image_removed)

    def legend_missing(prs):
        for shape in list(prs.slides[0].shapes):
            if getattr(shape, "has_text_frame", False) and any(token in shape.text for token in ("Legend", "Swimlanes:", "Boxes:", "Connectors:", "Badges:")):
                shape._element.getparent().remove(shape._element)
    save_mutation(args.golden, output_path(args.outdir, "legend-missing"), legend_missing)

    placeholder = args.outdir / "flatten-placeholder.png"
    Image.new("RGB", (1200, 900), "white").save(placeholder)
    def flattened(prs):
        slide = prs.slides[0]
        for shape in list(slide.shapes):
            shape._element.getparent().remove(shape._element)
        slide.shapes.add_picture(str(placeholder), 0, 0, width=prs.slide_width, height=prs.slide_height)
    save_mutation(args.golden, output_path(args.outdir, "flattened"), flattened)

    def hardcoded(prs):
        box = prs.slides[0].shapes.add_textbox(Inches(0.1), Inches(0.1), Inches(2.0), Inches(0.2))
        box.text = "solution/Golden"
    save_mutation(args.golden, output_path(args.outdir, "hardcoded-marker"), hardcoded)

    extra = args.outdir / "extra-output"
    extra.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.golden, extra / OUTPUT_NAME)
    (extra / "debug.txt").write_text("debug\n", encoding="utf-8")


if __name__ == "__main__":
    main()
