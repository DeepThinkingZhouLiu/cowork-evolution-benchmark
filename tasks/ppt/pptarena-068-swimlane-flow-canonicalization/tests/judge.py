#!/usr/bin/env python3
"""PPTArena Case 68 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import posixpath
import re
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE


OUTPUT_NAME = "SwimlaneFlowCanonicalization_GroundTruthA.pptx"
INPUT_NAME = "SwimlaneFlowCanonicalization_TestA.pptx"
INPUT_SHA256 = "c5f0e852c29a8347cb29de3897d2c16cb11c84c5a9cd32dfec0e69444c756982"
IMAGE_SHA256 = "75db072211fb94379043cfde485a8f4347fa15c360a6f8bae1835666add2be6e"
TASK_ID = "cowork-evo/pptarena-068-swimlane-flow-canonicalization"
RUBRIC_VERSION = "pptarena-case-068-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-068-deterministic-v0.1"
TITLE = "Canonical Process • 3 Swimlanes"
STEP_TEXTS = {index: f"Step {index}: {name}" for index, name in enumerate(("Intake", "Validate", "Transform", "Reject", "Export"), 1)}
STEP_FILLS = {1: "DCEBFF", 2: "E6FFE6", 3: "FFF0DC", 4: "FFE1E1", 5: "EBE1FF"}
FLOW = ((1, 2), (2, 3), (3, "decision"), ("decision", 4), ("decision", 5))
SPECS = {
    "R001": ("PPTX 完整且保留单页 4:3 画布", 6, "hurdle"),
    "R002": ("标题和 Intake/Processing/Output 三泳道正确", 15, "hurdle"),
    "R003": ("步骤、判定和图片归属、文字、尺寸与配色正确", 19, "hurdle"),
    "R004": ("1–5 圆形徽章与步骤一一对应", 13, "hurdle"),
    "R005": ("5 条折线连接符的关系和末端箭头正确", 20, "hurdle"),
    "R006": ("图例完整且重点元素无重叠越界", 12, "hurdle"),
    "R007": ("原生可编辑结构与原图片保留", 9, "hurdle"),
    "R008": ("可渲染且无宏、坏关系、答案标记或多余输出", 6, "positive"),
    "R009": ("无归属、徽章、连接、内容、图片、重叠或压平严重错误", -20, "penalty"),
}
SUSPICIOUS_MARKERS = ["solution/golden", "judge-result", "artifact_dir", "pptarena-case-068", "/root/"]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def safe(callback, default=""):
    try:
        return callback()
    except Exception:
        return default


def locate_source() -> Path | None:
    candidates = [Path("/app/input_files") / INPUT_NAME, Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME]
    return next((path for path in candidates if path.is_file()), None)


def package_checks(path: Path) -> tuple[bool, dict]:
    evidence = {"broken_member": None, "missing_targets": [], "macros": []}
    try:
        with zipfile.ZipFile(path) as archive:
            evidence["broken_member"] = archive.testzip()
            names = set(archive.namelist())
            evidence["macros"] = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
            for rels_name in sorted(name for name in names if name.endswith(".rels")):
                root = etree.fromstring(archive.read(rels_name))
                if rels_name == "_rels/.rels":
                    owner = ""
                else:
                    prefix, filename = rels_name.split("/_rels/", 1)
                    owner = posixpath.join(prefix, filename[:-5])
                base = posixpath.dirname(owner)
                for relation in root:
                    if relation.attrib.get("TargetMode") == "External":
                        continue
                    target = relation.attrib.get("Target", "")
                    resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                    if resolved not in names:
                        evidence["missing_targets"].append(f"{rels_name} -> {resolved}")
        return evidence["broken_member"] is None and not evidence["missing_targets"] and not evidence["macros"], evidence
    except Exception as exc:
        evidence["exception"] = str(exc)
        return False, evidence


def shape_text_map(slide) -> dict[str, list]:
    mapping: dict[str, list] = {}
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            mapping.setdefault(shape.text.strip(), []).append(shape)
    return mapping


def box(shape) -> tuple[int, int, int, int]:
    return shape.left, shape.top, shape.left + shape.width, shape.top + shape.height


def center(shape) -> tuple[float, float]:
    return shape.left + shape.width / 2, shape.top + shape.height / 2


def inside_slide(shape, presentation: Presentation) -> bool:
    left, top, right, bottom = box(shape)
    return left >= 0 and top >= 0 and right <= presentation.slide_width and bottom <= presentation.slide_height


def overlap_ratio(first, second) -> float:
    a = box(first)
    b = box(second)
    width = max(0, min(a[2], b[2]) - max(a[0], b[0]))
    height = max(0, min(a[3], b[3]) - max(a[1], b[1]))
    intersection = width * height
    denominator = min(first.width * first.height, second.width * second.height)
    return intersection / denominator if denominator else 0.0


def fill_rgb(shape) -> str:
    return safe(lambda: str(shape.fill.fore_color.rgb))


def lane_analysis(presentation: Presentation, text_map: dict[str, list]) -> tuple[bool, list, dict]:
    slide = presentation.slides[0]
    lanes = [
        shape
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
        and safe(lambda: shape.auto_shape_type == MSO_AUTO_SHAPE_TYPE.RECTANGLE, False)
        and shape.width >= presentation.slide_width * 0.8
        and presentation.slide_height * 0.17 <= shape.height <= presentation.slide_height * 0.35
    ]
    lanes.sort(key=lambda shape: shape.top)
    labels = ("Intake", "Processing", "Output")
    details = {"lane_ids": [shape.shape_id for shape in lanes], "labels": {}}
    passed = len(lanes) == 3
    if len(lanes) == 3:
        passed = passed and all(lanes[index].top + lanes[index].height <= lanes[index + 1].top + 91440 for index in range(2))
        passed = passed and max(lane.left for lane in lanes) - min(lane.left for lane in lanes) <= 91440
        passed = passed and max(lane.width for lane in lanes) - min(lane.width for lane in lanes) <= 91440
    for index, label in enumerate(labels):
        shapes = text_map.get(label, [])
        label_ok = len(shapes) == 1 and len(lanes) == 3
        if label_ok:
            x, y = center(shapes[0])
            label_ok = lanes[index].left <= x <= lanes[index].left + lanes[index].width and lanes[index].top <= y <= lanes[index].top + lanes[index].height
        details["labels"][label] = {"count": len(shapes), "ok": label_ok}
        passed = passed and label_ok
    return passed, lanes, details


def unique_text_shape(text_map: dict[str, list], text: str):
    shapes = text_map.get(text, [])
    return shapes[0] if len(shapes) == 1 else None


def process_analysis(presentation: Presentation, text_map: dict[str, list], lanes: list) -> tuple[bool, dict[int, object], object | None, object | None, dict]:
    steps = {index: unique_text_shape(text_map, text) for index, text in STEP_TEXTS.items()}
    decision = unique_text_shape(text_map, "OK?")
    pictures = [shape for shape in presentation.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    picture = pictures[0] if len(pictures) == 1 else None
    detail = {"steps": {}, "decision": None, "pictures": [shape.shape_id for shape in pictures]}
    passed = len(lanes) == 3 and all(shape is not None for shape in steps.values()) and decision is not None and picture is not None
    if all(shape is not None for shape in steps.values()):
        widths = [shape.width for shape in steps.values()]
        heights = [shape.height for shape in steps.values()]
        passed = passed and max(widths) - min(widths) <= max(widths) * 0.01 and max(heights) - min(heights) <= max(heights) * 0.01
    expected_lane = {1: 0, 2: 0, 3: 0, 4: 2, 5: 2}
    for index, shape in steps.items():
        current_ok = shape is not None and shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and safe(lambda: shape.auto_shape_type == MSO_AUTO_SHAPE_TYPE.RECTANGLE, False)
        if current_ok and len(lanes) == 3:
            _, y = center(shape)
            lane = lanes[expected_lane[index]]
            current_ok = lane.top <= y <= lane.top + lane.height and fill_rgb(shape) == STEP_FILLS[index] and inside_slide(shape, presentation)
        detail["steps"][index] = {"id": shape.shape_id if shape else None, "ok": current_ok, "fill": fill_rgb(shape) if shape else None}
        passed = passed and current_ok
    decision_ok = decision is not None and decision.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and safe(lambda: decision.auto_shape_type == MSO_AUTO_SHAPE_TYPE.DIAMOND, False)
    if decision_ok and len(lanes) == 3:
        _, y = center(decision)
        decision_ok = lanes[1].top <= y <= lanes[1].top + lanes[1].height and inside_slide(decision, presentation)
    detail["decision"] = {"id": decision.shape_id if decision else None, "ok": decision_ok}
    passed = passed and decision_ok
    picture_ok = picture is not None and len(lanes) == 3
    if picture_ok:
        _, y = center(picture)
        picture_ok = lanes[1].top <= y <= lanes[1].top + lanes[1].height and inside_slide(picture, presentation)
    detail["picture_ok"] = picture_ok
    passed = passed and picture_ok
    return passed, steps, decision, picture, detail


def badge_analysis(presentation: Presentation, steps: dict[int, object]) -> tuple[bool, dict]:
    slide = presentation.slides[0]
    badges = [
        shape
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
        and safe(lambda: shape.auto_shape_type == MSO_AUTO_SHAPE_TYPE.OVAL, False)
        and getattr(shape, "has_text_frame", False)
        and shape.text.strip() in {"1", "2", "3", "4", "5"}
    ]
    details = {"badges": []}
    passed = len(badges) == 5 and Counter(shape.text.strip() for shape in badges) == Counter({str(index): 1 for index in range(1, 6)})
    for index, step in steps.items():
        matches = [shape for shape in badges if shape.text.strip() == str(index)]
        associated = False
        if step is not None and len(matches) == 1:
            x, y = center(matches[0])
            padding = 182880
            associated = step.left - padding <= x <= step.left + step.width + padding and step.top - padding <= y <= step.top + step.height + padding
            associated = associated and matches[0].width > 0 and abs(matches[0].width - matches[0].height) <= max(matches[0].width, matches[0].height) * 0.08
        details["badges"].append({"number": index, "ids": [shape.shape_id for shape in matches], "associated": associated})
        passed = passed and associated
    return passed, details


def connector_endpoints(shape) -> tuple[int | None, int | None]:
    starts = shape._element.xpath(".//*[local-name()='stCxn']/@id")
    ends = shape._element.xpath(".//*[local-name()='endCxn']/@id")
    return (int(starts[0]) if len(starts) == 1 else None, int(ends[0]) if len(ends) == 1 else None)


def connector_geometry(shape) -> str:
    values = shape._element.xpath("./*[local-name()='spPr']/*[local-name()='prstGeom']/@prst")
    return values[0] if len(values) == 1 else ""


def connector_arrow(shape) -> str:
    values = shape._element.xpath("./*[local-name()='spPr']/*[local-name()='ln']/*[local-name()='tailEnd']/@type")
    return values[0] if len(values) == 1 else ""


def connector_style(shape) -> tuple:
    line = shape._element.xpath("./*[local-name()='spPr']/*[local-name()='ln']")
    if len(line) != 1:
        return ()
    width = line[0].attrib.get("w", "")
    colors = line[0].xpath("./*[local-name()='solidFill']/*/@val")
    return width, tuple(colors), connector_arrow(shape)


def connector_analysis(presentation: Presentation, steps: dict[int, object], decision) -> tuple[bool, dict]:
    connectors = [shape for shape in presentation.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.LINE]
    object_ids = {index: shape.shape_id if shape is not None else None for index, shape in steps.items()}
    object_ids["decision"] = decision.shape_id if decision is not None else None
    expected = {(object_ids[start], object_ids[end]) for start, end in FLOW}
    observed = {connector_endpoints(shape) for shape in connectors}
    styles = [connector_style(shape) for shape in connectors]
    details = {
        "count": len(connectors),
        "expected": sorted((str(a), str(b)) for a, b in expected),
        "observed": sorted((str(a), str(b)) for a, b in observed),
        "connectors": [{"id": shape.shape_id, "geometry": connector_geometry(shape), "arrow": connector_arrow(shape), "endpoints": connector_endpoints(shape), "style": connector_style(shape)} for shape in connectors],
    }
    passed = len(connectors) == 5 and None not in {value for pair in expected for value in pair}
    passed = passed and observed == expected
    passed = passed and all(connector_geometry(shape).startswith("bentConnector") for shape in connectors)
    passed = passed and all(connector_arrow(shape) not in ("", "none") for shape in connectors)
    passed = passed and len(set(styles)) == 1
    return passed, details


def legend_and_layout(presentation: Presentation, text_map: dict[str, list], steps: dict[int, object], decision, picture) -> tuple[bool, dict]:
    slide = presentation.slides[0]
    combined = "\n".join(shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)).casefold()
    required = ("legend", "swimlane", "intake", "processing", "output", "boxes", "connector", "elbow", "arrow", "badge", "1", "5")
    tokens_ok = all(token in combined for token in required)
    legend_shapes = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False) and any(token in shape.text.casefold() for token in ("legend", "swimlanes:", "boxes:", "connectors:", "badges:"))]
    important = [shape for shape in [*steps.values(), decision, picture] if shape is not None]
    overlap_pairs = []
    for index, first in enumerate(important):
        for second in important[index + 1 :]:
            ratio = overlap_ratio(first, second)
            if ratio > 0.04:
                overlap_pairs.append((first.shape_id, second.shape_id, round(ratio, 4)))
    legend_overlaps = []
    for legend in legend_shapes:
        for shape in important:
            ratio = overlap_ratio(legend, shape)
            if ratio > 0.04:
                legend_overlaps.append((legend.shape_id, shape.shape_id, round(ratio, 4)))
    outside = [shape.shape_id for shape in slide.shapes if shape.shape_id != 1 and not inside_slide(shape, presentation)]
    passed = tokens_ok and bool(legend_shapes) and not overlap_pairs and not legend_overlaps and not outside
    return passed, {"tokens_ok": tokens_ok, "legend_ids": [shape.shape_id for shape in legend_shapes], "important_overlaps": overlap_pairs, "legend_overlaps": legend_overlaps, "outside": outside}


def media_hashes(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as archive:
        return [hashlib.sha256(archive.read(name)).hexdigest() for name in archive.namelist() if name.startswith("ppt/media/") and not name.endswith("/")]


def render(path: Path, outdir: Path) -> tuple[int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "profile"
    profile.mkdir()
    completed = subprocess.run(
        ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(path)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=300,
        check=False,
    )
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return 0, completed.stdout[-500:]
    info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=30, check=False).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    pages = int(match.group(1)) if match else 0
    return pages, completed.stdout[-500:]


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {"criterion_id": identifier, "score": 1.0 if passed else 0.0, "evidence": evidence, "method": "deterministic", "raw": {"description": description, "weight": weight, "type": kind, "passed": bool(passed)}}


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [item["criterion_id"] for item in criteria if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0 and item["raw"]["passed"])
    penalty = sum(abs(item["raw"]["weight"]) for item in criteria if item["raw"]["weight"] < 0 and not item["raw"]["passed"])
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [sha256(path)] if path.is_file() else [], "criterion_results": criteria, "reward": reward, "passed": reward >= 0.8 and not hurdles, "hurdle_failures": hurdles}


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        return make_result(path, [result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配") for identifier in SPECS])
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        return make_result(path, [result_item(identifier, False, f"PPTX 打开失败：{exc}") for identifier in SPECS])

    criteria = []
    package_ok, package_evidence = package_checks(path)
    size_ok = len(source.slides) == len(candidate.slides) == 1 and (source.slide_width, source.slide_height) == (candidate.slide_width, candidate.slide_height)
    criteria.append(result_item("R001", package_ok and size_ok, f"页数={len(candidate.slides)}，尺寸一致={size_ok}，包={json.dumps(package_evidence, ensure_ascii=False)}"))

    text_map = shape_text_map(candidate.slides[0])
    title_ok = len(text_map.get(TITLE, [])) == 1
    lane_ok, lanes, lane_detail = lane_analysis(candidate, text_map)
    criteria.append(result_item("R002", title_ok and lane_ok, f"标题={title_ok}，{json.dumps(lane_detail, ensure_ascii=False)}"))

    process_ok, steps, decision, picture, process_detail = process_analysis(candidate, text_map, lanes)
    criteria.append(result_item("R003", process_ok, json.dumps(process_detail, ensure_ascii=False)))

    badge_ok, badge_detail = badge_analysis(candidate, steps)
    criteria.append(result_item("R004", badge_ok, json.dumps(badge_detail, ensure_ascii=False)))

    connector_ok, connector_detail = connector_analysis(candidate, steps, decision)
    criteria.append(result_item("R005", connector_ok, json.dumps(connector_detail, ensure_ascii=False)))

    legend_ok, layout_detail = legend_and_layout(candidate, text_map, steps, decision, picture)
    criteria.append(result_item("R006", legend_ok, json.dumps(layout_detail, ensure_ascii=False)))

    combined_text = "\n".join(shape.text for shape in candidate.slides[0].shapes if getattr(shape, "has_text_frame", False)).lower()
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker in combined_text]
    media = safe(lambda: media_hashes(path), [])
    full_slide_pictures = [shape.shape_id for shape in candidate.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.width * shape.height >= candidate.slide_width * candidate.slide_height * 0.8]
    editable_ok = len(media) == 1 and media == [IMAGE_SHA256] and len([shape for shape in candidate.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]) == 1 and not full_slide_pictures
    editable_ok = editable_ok and len([shape for shape in candidate.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.LINE]) == 5 and not markers
    criteria.append(result_item("R007", editable_ok, f"媒体={media}，整页图={full_slide_pictures}，答案标记={markers}"))

    render_pages = 0
    render_log = ""
    try:
        with tempfile.TemporaryDirectory(prefix="pptarena068-render-") as temporary:
            render_pages, render_log = render(path, Path(temporary))
    except Exception as exc:
        render_log = f"渲染异常：{exc}"
    delivered = sorted(str(item.relative_to(artifact_dir)) for item in artifact_dir.rglob("*") if item.is_file())
    hygiene_ok = render_pages == 1 and delivered == [OUTPUT_NAME] and package_ok and not markers
    criteria.append(result_item("R008", hygiene_ok, f"渲染页数={render_pages}，文件={delivered}，标记={markers}，日志={render_log}"))

    severe = not (title_ok and lane_ok and process_ok and badge_ok and connector_ok and legend_ok and editable_ok)
    criteria.append(result_item("R009", not severe, f"严重错误={severe}"))
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    path = args.artifact_dir / OUTPUT_NAME
    result = evaluate(path, args.artifact_dir) if path.is_file() else make_result(path, [result_item(identifier, False, f"缺失指定交付物：{OUTPUT_NAME}") for identifier in SPECS])
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{result['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
