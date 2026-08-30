#!/usr/bin/env python3
"""PPTArena Case 71 专用 Golden 构建器：逐单元格标记 6 处负数。"""

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor


INPUT_SHA256 = "0b64628a3f1169dc6dcfca2ea0acca80e1cc6aeba7b9aa9bc456e8f36b9a305f"
TARGETS = {
    (2, 2, 3): "-$15,250",
    (2, 4, 3): "-$15,500",
    (3, 2, 3): "−5.6%",
    (3, 3, 3): "(1.2%)",
    (4, 1, 1): "-12",
    (4, 4, 1): "-3",
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def only_table(presentation: Presentation, slide_number: int):
    tables = [shape.table for shape in presentation.slides[slide_number - 1].shapes if shape.has_table]
    if len(tables) != 1:
        raise ValueError(f"第 {slide_number} 页不是唯一表格")
    return tables[0]


def format_token(cell, token: str, digits_only: bool = False) -> None:
    if len(cell.text_frame.paragraphs) != 1:
        raise ValueError("目标单元格不再是单段落")
    paragraph = cell.text_frame.paragraphs[0]
    text = paragraph.text
    if text.count(token) != 1:
        raise ValueError(f"目标 {token!r} 计数异常：{text!r}")
    start = text.index(token)
    end = start + len(token)
    segments: list[tuple[str, bool, object]] = []
    cursor = 0
    for source_run in paragraph.runs:
        for character in source_run.text:
            selected = start <= cursor < end and (character.isdigit() if digits_only else True)
            if not segments or segments[-1][1] != selected or segments[-1][2] is not source_run:
                segments.append((character, selected, source_run))
            else:
                segments[-1] = (segments[-1][0] + character, selected, source_run)
            cursor += 1
    if cursor != len(text):
        raise AssertionError("分段长度错误")
    paragraph.clear()
    for segment, selected, source_run in segments:
        run = paragraph.add_run()
        run.text = segment
        source_font = source_run.font
        target_font = run.font
        if source_font.name is not None:
            target_font.name = source_font.name
        if source_font.size is not None:
            target_font.size = source_font.size
        if source_font.bold is not None:
            target_font.bold = source_font.bold
        if source_font.italic is not None:
            target_font.italic = source_font.italic
        if source_font.underline is not None:
            target_font.underline = source_font.underline
        if source_font.language_id is not None:
            target_font.language_id = source_font.language_id
        if source_font.color.type is not None:
            if source_font.color.rgb is not None:
                target_font.color.rgb = source_font.color.rgb
            elif source_font.color.theme_color is not None:
                target_font.color.theme_color = source_font.color.theme_color
                target_font.color.brightness = source_font.color.brightness
        if selected:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 0, 0)


def structure_signature(presentation: Presentation) -> tuple:
    return tuple(
        tuple((shape.shape_id, str(shape.shape_type), shape.left, shape.top, shape.width, shape.height) for shape in slide.shapes)
        for slide in presentation.slides
    )


def build(source: Path, output: Path, digits_only: bool = False) -> None:
    if sha256(source) != INPUT_SHA256:
        raise ValueError("冻结输入 SHA256 不匹配")
    original = Presentation(source)
    presentation = Presentation(source)
    if len(presentation.slides) != 4:
        raise ValueError("冻结输入不再是 4 页")
    for (slide_number, row, column), token in TARGETS.items():
        format_token(only_table(presentation, slide_number).cell(row, column), token, digits_only=digits_only)
    if structure_signature(presentation) != structure_signature(original):
        raise ValueError("对象结构或几何发生变化")
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    reopened = Presentation(output)
    if structure_signature(reopened) != structure_signature(original):
        raise ValueError("保存后对象结构或几何发生变化")
    for (slide_number, row, column), token in TARGETS.items():
        if token not in only_table(reopened, slide_number).cell(row, column).text:
            raise ValueError(f"保存后丢失负数：{token}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--digits-only", action="store_true")
    args = parser.parse_args()
    build(args.source, args.output, digits_only=args.digits_only)
    print(f"{args.output}\nsha256={sha256(args.output)}")


if __name__ == "__main__":
    main()
