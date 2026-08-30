#!/usr/bin/env python3
"""PPTArena Case 71 专用 Judge 校准样例。"""

from __future__ import annotations

import argparse
import importlib.util
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches


OUTPUT_NAME = "ConditionalTableFormatting_Negatives_GroundTruthA.pptx"


def output_path(root: Path, case: str) -> Path:
    path = root / case / OUTPUT_NAME
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def only_table(presentation: Presentation, slide_number: int):
    shapes = [shape for shape in presentation.slides[slide_number - 1].shapes if shape.has_table]
    if len(shapes) != 1:
        raise ValueError("缺少唯一表格")
    return shapes[0]


def target_runs(presentation: Presentation, slide_number: int, row: int, column: int):
    cell = only_table(presentation, slide_number).table.cell(row, column)
    return [run for paragraph in cell.text_frame.paragraphs for run in paragraph.runs if any(character.isdigit() for character in run.text)]


def save_mutation(golden: Path, output: Path, callback) -> None:
    presentation = Presentation(golden)
    callback(presentation)
    presentation.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--builder", type=Path, required=True)
    parser.add_argument("--slide2-image", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()
    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))

    specification = importlib.util.spec_from_file_location("case071_builder", args.builder)
    module = importlib.util.module_from_spec(specification)
    assert specification.loader is not None
    specification.loader.exec_module(module)
    module.build(args.source, output_path(args.outdir, "valid-digits-only"), digits_only=True)

    skipped = Presentation(args.source)
    for key, token in module.TARGETS.items():
        if key != (4, 4, 1):
            module.format_token(module.only_table(skipped, key[0]).cell(key[1], key[2]), token)
    skipped.save(output_path(args.outdir, "skip-one"))
    save_mutation(args.golden, output_path(args.outdir, "red-not-bold"), lambda prs: [setattr(run.font, "bold", False) for run in target_runs(prs, 3, 2, 3)])
    save_mutation(args.golden, output_path(args.outdir, "bold-black"), lambda prs: [setattr(run.font.color, "rgb", RGBColor(0, 0, 0)) for run in target_runs(prs, 2, 2, 3)])

    def positive(prs):
        run = only_table(prs, 3).table.cell(1, 1).text_frame.paragraphs[0].runs[0]
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 0, 0)
    save_mutation(args.golden, output_path(args.outdir, "format-positive"), positive)

    def all_cells(prs):
        for slide_number in (2, 3, 4):
            for row in only_table(prs, slide_number).table.rows:
                for cell in row.cells:
                    for paragraph in cell.text_frame.paragraphs:
                        for run in paragraph.runs:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 0, 0)
    save_mutation(args.golden, output_path(args.outdir, "format-all"), all_cells)

    def changed_text(prs):
        only_table(prs, 4).table.cell(1, 1).text_frame.paragraphs[0].runs[-1].text += "!"
    save_mutation(args.golden, output_path(args.outdir, "changed-text"), changed_text)

    def moved_table(prs):
        only_table(prs, 3).left += Inches(0.1)
    save_mutation(args.golden, output_path(args.outdir, "moved-table"), moved_table)

    def flattened(prs):
        slide = prs.slides[1]
        for shape in list(slide.shapes):
            shape._element.getparent().remove(shape._element)
        slide.shapes.add_picture(str(args.slide2_image), 0, 0, width=prs.slide_width, height=prs.slide_height)
    save_mutation(args.golden, output_path(args.outdir, "flattened"), flattened)

    def hardcoded(prs):
        box = prs.slides[0].shapes.add_textbox(Inches(0.1), Inches(0.1), Inches(2.0), Inches(0.2))
        box.text = "solution/Golden"
    save_mutation(args.golden, output_path(args.outdir, "hardcoded-marker"), hardcoded)

    extra = args.outdir / "extra-output"
    extra.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.golden, extra / OUTPUT_NAME)
    (extra / "debug.txt").write_text("debug\n", encoding="utf-8")


if __name__ == "__main__":
    main()
