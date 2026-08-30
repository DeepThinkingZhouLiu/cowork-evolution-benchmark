#!/usr/bin/env python3
"""PPTArena Case 71 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import math
import posixpath
import re
import statistics
import subprocess
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image, ImageChops, ImageDraw, ImageStat
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "ConditionalTableFormatting_Negatives_GroundTruthA.pptx"
INPUT_NAME = "ConditionalTableFormatting_Negatives_TestA.pptx"
INPUT_SHA256 = "0b64628a3f1169dc6dcfca2ea0acca80e1cc6aeba7b9aa9bc456e8f36b9a305f"
TASK_ID = "cowork-evo/pptarena-071-negative-number-table-formatting"
RUBRIC_VERSION = "pptarena-case-071-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-071-deterministic-v0.1"
SUSPICIOUS_MARKERS = ["solution/golden", "judge-result", "artifact_dir", "pptarena-case-071", "/root/"]
TARGETS = {
    (2, 2, 3): "-$15,250",
    (2, 4, 3): "-$15,500",
    (3, 2, 3): "−5.6%",
    (3, 3, 3): "(1.2%)",
    (4, 1, 1): "-12",
    (4, 4, 1): "-3",
}
SPECS = {
    "R001": ("PPTX 完整且严格保留原尺寸 4 页", 7, "hurdle"),
    "R002": ("6 处负数的数字均为可见红色粗体", 25, "hurdle"),
    "R003": ("表格文字、数据与负数表达完整保留", 18, "hurdle"),
    "R004": ("正数、非数字和负数外上下文样式保持", 15, "hurdle"),
    "R005": ("3 张原生表格的行列、几何、表格样式和可编辑性保持", 15, "hurdle"),
    "R006": ("封面、顺序、非表格对象、备注与全局结构保持", 10, "hurdle"),
    "R007": ("4 页可渲染，红色负数可见且目标单元格外区域保持", 6, "hurdle"),
    "R008": ("无宏、损坏关系、多余输出或答案标记", 4, "positive"),
    "R009": ("无漏改、错样式、数据改写、正数误改或压平等严重错误", -20, "penalty"),
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def safe(callback, default=""):
    try:
        return callback()
    except Exception:
        return default


def locate_source() -> Path | None:
    candidates = [Path("/app/input_files") / INPUT_NAME, Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME]
    return next((path for path in candidates if path.is_file()), None)


def only_table(presentation: Presentation, slide_number: int):
    shapes = [shape for shape in presentation.slides[slide_number - 1].shapes if shape.has_table]
    if len(shapes) != 1:
        raise ValueError(f"第 {slide_number} 页表格数不是 1")
    return shapes[0]


def font_signature(font) -> tuple:
    color = font.color
    return (
        font.name,
        round(font.size.pt, 4) if font.size else None,
        font.bold,
        font.italic,
        font.underline,
        safe(lambda: str(font.language_id)),
        safe(lambda: str(color.type)),
        safe(lambda: str(color.rgb)),
        safe(lambda: str(color.theme_color)),
        safe(lambda: round(color.brightness, 5)),
    )


def character_styles(paragraph) -> tuple:
    return tuple(style for run in paragraph.runs for style in [font_signature(run.font)] * len(run.text))


def canonical(element) -> bytes:
    if element is None:
        return b""
    return etree.tostring(element, method="c14n")


def paragraph_format(paragraph) -> bytes:
    return canonical(paragraph._p.pPr)


def frame_signature(frame) -> tuple:
    return (
        frame.margin_left,
        frame.margin_right,
        frame.margin_top,
        frame.margin_bottom,
        str(frame.vertical_anchor),
        frame.word_wrap,
        str(frame.auto_size),
    )


def cell_signature(cell) -> tuple:
    return (
        cell.text,
        frame_signature(cell.text_frame),
        canonical(cell._tc.tcPr),
        tuple((paragraph.text, character_styles(paragraph), paragraph_format(paragraph)) for paragraph in cell.text_frame.paragraphs),
    )


def non_text_table_signature(shape) -> bytes:
    element = copy.deepcopy(shape._element)
    for body in element.xpath(".//*[local-name()='txBody']"):
        for paragraph in list(body.xpath("./*[local-name()='p']")):
            body.remove(paragraph)
    return canonical(element)


def shape_structure(presentation: Presentation) -> tuple:
    return tuple(
        tuple(
            (
                shape.shape_id,
                str(shape.shape_type),
                shape.left,
                shape.top,
                shape.width,
                shape.height,
                str(shape.placeholder_format.type) if shape.is_placeholder else "",
            )
            for shape in slide.shapes
        )
        for slide in presentation.slides
    )


def text_shape_signature(shape) -> tuple:
    return (
        shape.text,
        frame_signature(shape.text_frame),
        tuple((paragraph.text, character_styles(paragraph), paragraph_format(paragraph)) for paragraph in shape.text_frame.paragraphs),
    )


def notes_signature(slide) -> tuple[str, ...]:
    try:
        return tuple(shape.text for shape in slide.notes_slide.shapes if getattr(shape, "has_text_frame", False))
    except Exception:
        return ()


def package_checks(path: Path) -> tuple[bool, dict]:
    evidence = {"broken_member": None, "missing_targets": [], "macros": []}
    try:
        with zipfile.ZipFile(path) as archive:
            evidence["broken_member"] = archive.testzip()
            names = set(archive.namelist())
            evidence["macros"] = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
            for rels_name in sorted(name for name in names if name.endswith(".rels")):
                root = etree.fromstring(archive.read(rels_name))
                if rels_name == "_rels/.rels":
                    owner = ""
                else:
                    prefix, filename = rels_name.split("/_rels/", 1)
                    owner = posixpath.join(prefix, filename[:-5])
                base = posixpath.dirname(owner)
                for relation in root:
                    if relation.attrib.get("TargetMode") == "External":
                        continue
                    target = relation.attrib.get("Target", "")
                    resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                    if resolved not in names:
                        evidence["missing_targets"].append(f"{rels_name} -> {resolved}")
        passed = evidence["broken_member"] is None and not evidence["missing_targets"] and not evidence["macros"]
        return passed, evidence
    except Exception as exc:
        evidence["exception"] = str(exc)
        return False, evidence


def is_red(style: tuple) -> bool:
    rgb = style[7]
    if not re.fullmatch(r"[0-9A-Fa-f]{6}", rgb or ""):
        return False
    red, green, blue = int(rgb[:2], 16), int(rgb[2:4], 16), int(rgb[4:], 16)
    return red >= 180 and green <= 70 and blue <= 70 and red >= green * 2.5 and red >= blue * 2.5


def target_cell_check(source_cell, candidate_cell, token: str) -> tuple[bool, bool, bool, dict]:
    source_paragraphs = source_cell.text_frame.paragraphs
    candidate_paragraphs = candidate_cell.text_frame.paragraphs
    text_ok = candidate_cell.text == source_cell.text and source_cell.text.count(token) == 1
    frame_ok = frame_signature(candidate_cell.text_frame) == frame_signature(source_cell.text_frame) and canonical(candidate_cell._tc.tcPr) == canonical(source_cell._tc.tcPr)
    if len(source_paragraphs) != 1 or len(candidate_paragraphs) != 1 or not text_ok:
        return False, text_ok, False, {"text": candidate_cell.text}
    source_paragraph = source_paragraphs[0]
    candidate_paragraph = candidate_paragraphs[0]
    source_styles = character_styles(source_paragraph)
    candidate_styles = character_styles(candidate_paragraph)
    if len(source_styles) != len(candidate_styles) or paragraph_format(source_paragraph) != paragraph_format(candidate_paragraph):
        return False, text_ok, False, {"style_lengths": [len(source_styles), len(candidate_styles)]}
    start = source_cell.text.index(token)
    end = start + len(token)
    target_ok = True
    preservation_ok = frame_ok
    digit_count = 0
    for index, (character, source_style, candidate_style) in enumerate(zip(source_cell.text, source_styles, candidate_styles, strict=True)):
        if start <= index < end:
            if source_style[:2] + source_style[3:6] != candidate_style[:2] + candidate_style[3:6]:
                preservation_ok = False
            if character.isdigit():
                digit_count += 1
                if candidate_style[2] is not True or not is_red(candidate_style):
                    target_ok = False
        elif source_style != candidate_style:
            preservation_ok = False
    return target_ok and digit_count > 0, text_ok, preservation_ok, {"token": token, "digits": digit_count, "frame": frame_ok}


def render(path: Path, outdir: Path) -> tuple[list[Path], int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "profile"
    profile.mkdir()
    completed = subprocess.run(
        ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(path)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=300,
        check=False,
    )
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return [], 0, completed.stdout[-500:]
    info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=30, check=False).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    pages = int(match.group(1)) if match else 0
    prefix = outdir / "slide"
    raster = subprocess.run(["pdftoppm", "-png", "-r", "100", str(pdf), str(prefix)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300, check=False)
    images = sorted(outdir.glob("slide-*.png")) if raster.returncode == 0 else []
    return images, pages, f"soffice={completed.returncode}, raster={raster.returncode}"


def cell_box(shape, row: int, column: int, image: Image.Image, presentation: Presentation) -> tuple[int, int, int, int]:
    left = shape.left + sum(shape.table.columns[index].width for index in range(column))
    top = shape.top + sum(shape.table.rows[index].height for index in range(row))
    right = left + shape.table.columns[column].width
    bottom = top + shape.table.rows[row].height
    return (
        max(0, int(left * image.width / presentation.slide_width) - 2),
        max(0, int(top * image.height / presentation.slide_height) - 2),
        min(image.width, int(math.ceil(right * image.width / presentation.slide_width)) + 2),
        min(image.height, int(math.ceil(bottom * image.height / presentation.slide_height)) + 2),
    )


def render_metrics(source_images: list[Path], candidate_images: list[Path], source: Presentation, candidate: Presentation) -> tuple[float, int]:
    similarities = []
    red_pixels = 0
    targets_by_slide: dict[int, list[tuple[int, int]]] = {}
    for slide_number, row, column in TARGETS:
        targets_by_slide.setdefault(slide_number, []).append((row, column))
    for slide_number, (source_path, candidate_path) in enumerate(zip(source_images, candidate_images, strict=True), 1):
        with Image.open(source_path) as source_raw, Image.open(candidate_path) as candidate_raw:
            first = source_raw.convert("RGB")
            second = candidate_raw.convert("RGB")
            if first.size != second.size:
                similarities.append(0.0)
                continue
            mask = Image.new("L", first.size, 255)
            draw = ImageDraw.Draw(mask)
            if slide_number in targets_by_slide:
                table_shape = only_table(candidate, slide_number)
                for row, column in targets_by_slide[slide_number]:
                    box = cell_box(table_shape, row, column, second, candidate)
                    draw.rectangle(box, fill=0)
                    crop = second.crop(box)
                    red_pixels += sum(1 for red, green, blue in crop.getdata() if red >= 180 and green <= 90 and blue <= 90 and red > green * 2 and red > blue * 2)
            difference = ImageChops.difference(first, second)
            difference.paste((0, 0, 0), mask=ImageChops.invert(mask))
            similarities.append(max(0.0, 1.0 - statistics.mean(ImageStat.Stat(difference).rms) / 255.0))
    return min(similarities, default=0.0), red_pixels


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {"criterion_id": identifier, "score": 1.0 if passed else 0.0, "evidence": evidence, "method": "deterministic", "raw": {"description": description, "weight": weight, "type": kind, "passed": bool(passed)}}


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [item["criterion_id"] for item in criteria if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0 and item["raw"]["passed"])
    penalty = sum(abs(item["raw"]["weight"]) for item in criteria if item["raw"]["weight"] < 0 and not item["raw"]["passed"])
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [sha256(path)] if path.is_file() else [], "criterion_results": criteria, "reward": reward, "passed": reward >= 0.8 and not hurdles, "hurdle_failures": hurdles}


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        return make_result(path, [result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配") for identifier in SPECS])
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        return make_result(path, [result_item(identifier, False, f"PPTX 打开失败：{exc}") for identifier in SPECS])

    criteria = []
    package_passed, package_evidence = package_checks(path)
    size_passed = len(source.slides) == len(candidate.slides) == 4 and (source.slide_width, source.slide_height) == (candidate.slide_width, candidate.slide_height)
    criteria.append(result_item("R001", package_passed and size_passed, f"页数={len(candidate.slides)}，尺寸一致={size_passed}，包={json.dumps(package_evidence, ensure_ascii=False)}"))

    target_passed = True
    text_passed = True
    target_preservation = True
    target_details = []
    try:
        for key, token in TARGETS.items():
            source_cell = only_table(source, key[0]).table.cell(key[1], key[2])
            candidate_cell = only_table(candidate, key[0]).table.cell(key[1], key[2])
            target_ok, text_ok, preserve_ok, detail = target_cell_check(source_cell, candidate_cell, token)
            target_passed = target_passed and target_ok
            text_passed = text_passed and text_ok
            target_preservation = target_preservation and preserve_ok
            target_details.append({"cell": key, "target": target_ok, "text": text_ok, "preserve": preserve_ok, **detail})
    except Exception as exc:
        target_passed = text_passed = target_preservation = False
        target_details.append({"exception": str(exc)})
    criteria.append(result_item("R002", target_passed, json.dumps(target_details, ensure_ascii=False)))

    all_table_text_passed = True
    other_styles_passed = target_preservation
    table_structure_passed = True
    table_counts = []
    try:
        for slide_number in (2, 3, 4):
            source_shape = only_table(source, slide_number)
            candidate_shape = only_table(candidate, slide_number)
            source_table = source_shape.table
            candidate_table = candidate_shape.table
            table_counts.append((slide_number, len(candidate_table.rows), len(candidate_table.columns)))
            table_structure_passed = table_structure_passed and non_text_table_signature(source_shape) == non_text_table_signature(candidate_shape)
            for row in range(len(source_table.rows)):
                for column in range(len(source_table.columns)):
                    source_cell = source_table.cell(row, column)
                    candidate_cell = candidate_table.cell(row, column)
                    all_table_text_passed = all_table_text_passed and source_cell.text == candidate_cell.text
                    if (slide_number, row, column) not in TARGETS:
                        other_styles_passed = other_styles_passed and cell_signature(source_cell) == cell_signature(candidate_cell)
    except Exception:
        all_table_text_passed = other_styles_passed = table_structure_passed = False
    criteria.append(result_item("R003", text_passed and all_table_text_passed, f"所有表格文字保持={all_table_text_passed}，目标文字保持={text_passed}"))
    criteria.append(result_item("R004", other_styles_passed, f"非目标单元格和目标外字符样式保持={other_styles_passed}"))

    structure_equal = shape_structure(source) == shape_structure(candidate)
    full_slide_pictures = [
        (slide_number, shape.shape_id)
        for slide_number, slide in enumerate(candidate.slides, 1)
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.width * shape.height >= candidate.slide_width * candidate.slide_height * 0.8
    ]
    editable_passed = structure_equal and table_structure_passed and table_counts == [(2, 5, 4), (3, 4, 5), (4, 5, 2)] and not full_slide_pictures
    criteria.append(result_item("R005", editable_passed, f"表格={table_counts}，结构={structure_equal}，表格非文本 XML={table_structure_passed}，整页图={full_slide_pictures}"))

    global_passed = structure_equal
    for slide_number, (source_slide, candidate_slide) in enumerate(zip(source.slides, candidate.slides, strict=False), 1):
        global_passed = global_passed and notes_signature(source_slide) == notes_signature(candidate_slide)
        for source_shape, candidate_shape in zip(source_slide.shapes, candidate_slide.shapes, strict=False):
            if source_shape.has_table:
                continue
            if getattr(source_shape, "has_text_frame", False):
                global_passed = global_passed and text_shape_signature(source_shape) == text_shape_signature(candidate_shape)
    criteria.append(result_item("R006", global_passed, f"对象与非表格内容保持={global_passed}"))

    render_passed = False
    render_evidence = "未渲染"
    try:
        with tempfile.TemporaryDirectory(prefix="pptarena071-render-") as temporary:
            root = Path(temporary)
            source_images, source_pages, source_log = render(source_path, root / "source")
            candidate_images, candidate_pages, candidate_log = render(path, root / "candidate")
            if source_pages == candidate_pages == 4 and len(source_images) == len(candidate_images) == 4:
                similarity, red_pixels = render_metrics(source_images, candidate_images, source, candidate)
                render_passed = similarity >= 0.999 and red_pixels >= 30
                render_evidence = f"目标单元格外最低相似度={similarity:.6f}，红色像素={red_pixels}"
            else:
                render_evidence = f"页数={source_pages}/{candidate_pages}，{source_log}，{candidate_log}"
    except Exception as exc:
        render_evidence = f"渲染异常：{exc}"
    criteria.append(result_item("R007", render_passed, render_evidence))

    delivered = sorted(str(item.relative_to(artifact_dir)) for item in artifact_dir.rglob("*") if item.is_file())
    combined_text = "\n".join(shape.text for slide in candidate.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False)).lower()
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker in combined_text]
    hygiene_passed = delivered == [OUTPUT_NAME] and package_passed and not markers
    criteria.append(result_item("R008", hygiene_passed, f"文件={delivered}，标记={markers}"))

    severe = not (target_passed and text_passed and all_table_text_passed and other_styles_passed and editable_passed and global_passed)
    criteria.append(result_item("R009", not severe, f"严重错误={severe}"))
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    path = args.artifact_dir / OUTPUT_NAME
    if path.is_file():
        result = evaluate(path, args.artifact_dir)
    else:
        result = make_result(path, [result_item(identifier, False, f"缺失指定交付物：{OUTPUT_NAME}") for identifier in SPECS])
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{result['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
