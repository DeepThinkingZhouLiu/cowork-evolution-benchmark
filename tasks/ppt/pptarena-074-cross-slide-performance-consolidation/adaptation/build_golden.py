#!/usr/bin/env python3
"""PPTArena Case 74 专用 Golden 精修器。"""

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt


INPUT_SHA256 = "76214e891a723b38a67f17012099427b66c201b066150a76e0de5ed7befe3c07"
UPSTREAM_SHA256 = "57da6f276ec25e91c6e6c27dbc7ad30b471d8c31dc6645a5e033b4fc0f8e37e3"
ROWS = (
    ("Category", "Team A", "Team B"),
    ("Manufacturing", "20", "30"),
    ("Finance", "25", "35"),
    ("IT", "40", "10"),
    ("Other", "Logistics – 35\nHR – 15", "Sales – 20\nMarketing – 25"),
)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def only_table(presentation: Presentation):
    tables = [shape for shape in presentation.slides[1].shapes if shape.has_table]
    if len(tables) != 1:
        raise ValueError("第 2 页不再只有一张表格")
    return tables[0]


def set_body_cell(cell, text: str) -> None:
    cell.text = text
    for paragraph in cell.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(11)


def build(source: Path, upstream: Path, output: Path) -> None:
    if sha256(source) != INPUT_SHA256 or sha256(upstream) != UPSTREAM_SHA256:
        raise ValueError("冻结输入或上游 Ground Truth SHA256 不匹配")
    presentation = Presentation(upstream)
    if len(presentation.slides) != 3:
        raise ValueError("上游 Ground Truth 不再是 3 页")
    table = only_table(presentation).table
    if len(table.rows) != 5 or len(table.columns) != 3:
        raise ValueError("上游汇总表不再是 5×3")
    for row_index, values in enumerate(ROWS):
        for column_index, value in enumerate(values):
            if row_index == 0:
                if table.cell(row_index, column_index).text != value:
                    raise ValueError("表头与已审计契约不匹配")
            else:
                set_body_cell(table.cell(row_index, column_index), value)
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    reopened = Presentation(output)
    observed = tuple(tuple(cell.text for cell in row.cells) for row in only_table(reopened).table.rows)
    if observed != ROWS:
        raise ValueError("汇总表保存后与已审计数据不匹配")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--upstream", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.source, args.upstream, args.output)
    print(f"{args.output}\nsha256={sha256(args.output)}")


if __name__ == "__main__":
    main()
