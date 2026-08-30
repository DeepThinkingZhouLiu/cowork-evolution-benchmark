#!/usr/bin/env python3
"""PPTArena Case 74 专用 Judge 校准样例。"""

from __future__ import annotations

import argparse
import shutil
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches


OUTPUT_NAME = "CrossSlideDataConsolidation_GroundTruthA.pptx"


def output_path(root: Path, case: str) -> Path:
    path = root / case / OUTPUT_NAME
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def save_mutation(golden: Path, output: Path, callback) -> None:
    presentation = Presentation(golden)
    callback(presentation)
    presentation.save(output)


def only_table(presentation: Presentation):
    matches = [shape for shape in presentation.slides[1].shapes if shape.has_table]
    if len(matches) != 1:
        raise ValueError("缺少唯一表格")
    return matches[0].table


def by_text(presentation: Presentation, text: str):
    matches = [shape for slide in presentation.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text == text]
    if len(matches) != 1:
        raise ValueError(f"文字对象不唯一：{text!r}")
    return matches[0]


def repackage(golden: Path, output: Path) -> None:
    with zipfile.ZipFile(golden) as reader, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_STORED) as writer:
        for info in reader.infolist():
            writer.writestr(info.filename, reader.read(info.filename))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--upstream", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()
    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))
    shutil.copy2(args.upstream, output_path(args.outdir, "upstream-valid"))
    repackage(args.golden, output_path(args.outdir, "valid-repackage"))

    save_mutation(args.golden, output_path(args.outdir, "wrong-common-score"), lambda prs: setattr(only_table(prs).cell(1, 1), "text", "99"))
    save_mutation(args.golden, output_path(args.outdir, "missing-unique-fact"), lambda prs: setattr(only_table(prs).cell(4, 1), "text", "Logistics – 35"))
    save_mutation(args.golden, output_path(args.outdir, "swapped-team"), lambda prs: setattr(only_table(prs).cell(4, 2), "text", "Logistics – 35\nHR – 15"))
    save_mutation(args.golden, output_path(args.outdir, "title-case-failure"), lambda prs: setattr(by_text(prs, "Images and Captions").text_frame.paragraphs[0].runs[0], "text", "images and captions"))

    def unequal_image(prs):
        pictures = [shape for shape in prs.slides[2].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        pictures[1].width += Inches(0.3)
    save_mutation(args.golden, output_path(args.outdir, "unequal-image"), unequal_image)

    def misaligned_image(prs):
        pictures = [shape for shape in prs.slides[2].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        pictures[2].top += Inches(0.4)
    save_mutation(args.golden, output_path(args.outdir, "misaligned-image"), misaligned_image)

    def left_caption(prs):
        by_text(prs, "Finance: Financial Growth").text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    save_mutation(args.golden, output_path(args.outdir, "left-caption"), left_caption)

    def image_removed(prs):
        picture = next(shape for shape in prs.slides[2].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
        picture._element.getparent().remove(picture._element)
    save_mutation(args.golden, output_path(args.outdir, "image-removed"), image_removed)

    placeholder = args.outdir / "flatten-placeholder.png"
    Image.new("RGB", (1200, 900), "white").save(placeholder)
    def flattened(prs):
        slide = prs.slides[2]
        for shape in list(slide.shapes):
            shape._element.getparent().remove(shape._element)
        slide.shapes.add_picture(str(placeholder), 0, 0, width=prs.slide_width, height=prs.slide_height)
    save_mutation(args.golden, output_path(args.outdir, "flattened-final-slide"), flattened)

    def hardcoded(prs):
        box = prs.slides[0].shapes.add_textbox(Inches(0.1), Inches(0.1), Inches(2.0), Inches(0.2))
        box.text = "solution/Golden"
    save_mutation(args.golden, output_path(args.outdir, "hardcoded-marker"), hardcoded)

    extra = args.outdir / "extra-output"
    extra.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.golden, extra / OUTPUT_NAME)
    (extra / "debug.txt").write_text("debug\n", encoding="utf-8")


if __name__ == "__main__":
    main()
