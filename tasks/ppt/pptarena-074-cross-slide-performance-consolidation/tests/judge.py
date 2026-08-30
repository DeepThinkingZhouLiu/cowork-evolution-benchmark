#!/usr/bin/env python3
"""PPTArena Case 74 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import posixpath
import re
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN


OUTPUT_NAME = "CrossSlideDataConsolidation_GroundTruthA.pptx"
INPUT_NAME = "CrossSlideDataConsolidation_TestA.pptx"
INPUT_SHA256 = "76214e891a723b38a67f17012099427b66c201b066150a76e0de5ed7befe3c07"
TASK_ID = "cowork-evo/pptarena-074-cross-slide-performance-consolidation"
RUBRIC_VERSION = "pptarena-case-074-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-074-deterministic-v0.1"
EXPECTED_FACTS = {
    ("Team A", "Manufacturing"): 20,
    ("Team A", "Logistics"): 35,
    ("Team A", "Finance"): 25,
    ("Team A", "HR"): 15,
    ("Team A", "IT"): 40,
    ("Team B", "Manufacturing"): 30,
    ("Team B", "Sales"): 20,
    ("Team B", "Finance"): 35,
    ("Team B", "Marketing"): 25,
    ("Team B", "IT"): 10,
}
COMMON = {"Manufacturing", "Finance", "IT"}
UNIQUE_A = {"Logistics", "HR"}
UNIQUE_B = {"Sales", "Marketing"}
IMAGE_ORDER = (
    "268afe5405c04c0fdc6f60c20c42455e6f4bcfa652624bd0dd7d6113add448be",
    "fb3b9cc59056870ca64b76b3813e19f5cf3fe5f1c273d39d71d61e02af4be322",
    "ad20f6b172f348e93282fc4de6d1d34a888ab99801374be2f8aef0755af418df",
)
CAPTIONS = (
    "Manufacturing: Assembly Line",
    "Finance: Financial Growth",
    "IT: Data Servers",
)
REQUIRED_TEXT = (
    "Project Performance Review",
    "Quarterly Analysis by Teams",
    "Summary Performance",
    "Images and Captions",
)
SPECS = {
    "R001": ("PPTX 完整、保留原画布且最终 3 页", 7, "hurdle"),
    "R002": ("Title Case 标题副标题正确且原两页已删除", 12, "hurdle"),
    "R003": ("3 列原生表格中 10 个分数事实完整准确", 25, "hurdle"),
    "R004": ("项目符号列表移除、表格为主内容且无图表残留", 12, "hurdle"),
    "R005": ("三张原图哈希、顺序、尺寸和顶部对齐正确", 16, "hurdle"),
    "R006": ("三条原图注与图片关联、居中与布局正确", 12, "hurdle"),
    "R007": ("封面与原媒体保留，无压平、整页图或答案标记", 10, "hurdle"),
    "R008": ("3 页可渲染且无宏、坏关系或多余输出", 6, "positive"),
    "R009": ("无页数、数据、列表、图表、图片、图注或压平严重错误", -20, "penalty"),
}
SUSPICIOUS_MARKERS = ["solution/golden", "judge-result", "artifact_dir", "pptarena-case-074", "/root/"]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def safe(callback, default=""):
    try:
        return callback()
    except Exception:
        return default


def locate_source() -> Path | None:
    candidates = [Path("/app/input_files") / INPUT_NAME, Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME]
    return next((path for path in candidates if path.is_file()), None)


def package_checks(path: Path) -> tuple[bool, dict]:
    evidence = {"broken_member": None, "missing_targets": [], "macros": []}
    try:
        with zipfile.ZipFile(path) as archive:
            evidence["broken_member"] = archive.testzip()
            names = set(archive.namelist())
            evidence["macros"] = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
            for rels_name in sorted(name for name in names if name.endswith(".rels")):
                root = etree.fromstring(archive.read(rels_name))
                if rels_name == "_rels/.rels":
                    owner = ""
                else:
                    prefix, filename = rels_name.split("/_rels/", 1)
                    owner = posixpath.join(prefix, filename[:-5])
                base = posixpath.dirname(owner)
                for relation in root:
                    if relation.attrib.get("TargetMode") == "External":
                        continue
                    target = relation.attrib.get("Target", "")
                    resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                    if resolved not in names:
                        evidence["missing_targets"].append(f"{rels_name} -> {resolved}")
        return evidence["broken_member"] is None and not evidence["missing_targets"] and not evidence["macros"], evidence
    except Exception as exc:
        evidence["exception"] = str(exc)
        return False, evidence


def all_text(presentation: Presentation) -> list[str]:
    return [shape.text for slide in presentation.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False)]


def only_table(presentation: Presentation):
    if len(presentation.slides) < 2:
        return None
    tables = [shape for shape in presentation.slides[1].shapes if shape.has_table]
    return tables[0] if len(tables) == 1 else None


def clean_category(text: str) -> str:
    normalized = re.sub(r"\s+", " ", text.replace("\v", " ").replace("\n", " ")).strip()
    for category in COMMON | UNIQUE_A | UNIQUE_B | {"Other"}:
        if normalized.casefold() == category.casefold():
            return category
    return normalized


def sole_integer(text: str) -> int | None:
    values = [int(value) for value in re.findall(r"(?<!\d)\d+(?!\d)", text)]
    return values[0] if len(values) == 1 else None


def named_pairs(text: str, allowed: set[str]) -> dict[str, int]:
    normalized = text.replace("\v", " ").replace("\n", " ")
    pairs = {}
    for category in allowed:
        match = re.search(rf"\b{re.escape(category)}\b\s*[-–—:]?\s*(\d+)", normalized, re.IGNORECASE)
        if match:
            pairs[category] = int(match.group(1))
    return pairs


def table_facts(table_shape) -> tuple[dict, list, dict]:
    table = table_shape.table
    observed: dict[tuple[str, str], int] = {}
    errors = []
    header = [clean_category(cell.text) for cell in table.rows[0].cells] if len(table.rows) else []
    for row_index in range(1, len(table.rows)):
        row = table.rows[row_index]
        if len(row.cells) != 3:
            errors.append(f"第 {row_index + 1} 行不是 3 列")
            continue
        category = clean_category(row.cells[0].text)
        team_a = row.cells[1].text
        team_b = row.cells[2].text
        if category in COMMON:
            observed[("Team A", category)] = sole_integer(team_a)
            observed[("Team B", category)] = sole_integer(team_b)
        elif category in UNIQUE_A:
            observed[("Team A", category)] = sole_integer(team_a)
            if sole_integer(team_b) is not None:
                errors.append(f"{category} 误写 Team B 分数")
        elif category in UNIQUE_B:
            observed[("Team B", category)] = sole_integer(team_b)
            if sole_integer(team_a) is not None:
                errors.append(f"{category} 误写 Team A 分数")
        elif category == "Other":
            for name, value in named_pairs(team_a, UNIQUE_A).items():
                observed[("Team A", name)] = value
            for name, value in named_pairs(team_b, UNIQUE_B).items():
                observed[("Team B", name)] = value
            wrong_a = named_pairs(team_a, UNIQUE_B)
            wrong_b = named_pairs(team_b, UNIQUE_A)
            if wrong_a or wrong_b:
                errors.append(f"Other 行串队：A={wrong_a}，B={wrong_b}")
        else:
            errors.append(f"未知类别：{category!r}")
    return observed, errors, {"header": header, "rows": len(table.rows), "columns": len(table.columns)}


def picture_hash(shape) -> str:
    return hashlib.sha256(shape.image.blob).hexdigest()


def overlap_ratio(first, second) -> float:
    left = max(first.left, second.left)
    top = max(first.top, second.top)
    right = min(first.left + first.width, second.left + second.width)
    bottom = min(first.top + first.height, second.top + second.height)
    intersection = max(0, right - left) * max(0, bottom - top)
    denominator = min(first.width * first.height, second.width * second.height)
    return intersection / denominator if denominator else 0.0


def final_slide_analysis(presentation: Presentation) -> tuple[bool, bool, dict]:
    if len(presentation.slides) != 3:
        return False, False, {"reason": "页数不是 3"}
    slide = presentation.slides[2]
    pictures = sorted([shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE], key=lambda shape: shape.left)
    captions = []
    for text in CAPTIONS:
        matches = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip() == text]
        captions.append(matches[0] if len(matches) == 1 else None)
    image_ok = len(pictures) == 3
    if image_ok:
        hashes = tuple(picture_hash(shape) for shape in pictures)
        dimensions = [(shape.width, shape.height) for shape in pictures]
        tops = [shape.top for shape in pictures]
        image_ok = hashes == IMAGE_ORDER
        image_ok = image_ok and max(width for width, _ in dimensions) - min(width for width, _ in dimensions) <= max(width for width, _ in dimensions) * 0.01
        image_ok = image_ok and max(height for _, height in dimensions) - min(height for _, height in dimensions) <= max(height for _, height in dimensions) * 0.01
        image_ok = image_ok and max(tops) - min(tops) <= 18288
    else:
        hashes, dimensions, tops = (), [], []
    caption_ok = all(shape is not None for shape in captions) and len(pictures) == 3
    associations = []
    if caption_ok:
        for picture, caption in zip(pictures, captions, strict=True):
            centered = all(paragraph.alignment == PP_ALIGN.CENTER for paragraph in caption.text_frame.paragraphs)
            horizontal = abs((picture.left + picture.width / 2) - (caption.left + caption.width / 2)) <= 91440
            below = caption.top >= picture.top + picture.height - 9144
            width_match = abs(caption.width - picture.width) <= max(caption.width, picture.width) * 0.01
            associations.append({"picture": picture.shape_id, "caption": caption.shape_id, "centered": centered, "horizontal": horizontal, "below": below, "width_match": width_match})
            caption_ok = caption_ok and centered and horizontal and below and width_match
        for index, picture in enumerate(pictures):
            for other in pictures[index + 1 :]:
                caption_ok = caption_ok and overlap_ratio(picture, other) == 0
    return image_ok, caption_ok, {"hashes": hashes, "dimensions": dimensions, "tops": tops, "associations": associations}


def source_cover_preserved(source: Presentation, candidate: Presentation) -> bool:
    if not candidate.slides:
        return False
    source_slide = source.slides[0]
    candidate_slide = candidate.slides[0]
    if len(source_slide.shapes) != len(candidate_slide.shapes) != 2:
        return False
    expected = ("Project Performance Review", "Quarterly Analysis by Teams")
    for index, (source_shape, candidate_shape) in enumerate(zip(source_slide.shapes, candidate_slide.shapes, strict=False)):
        if (source_shape.left, source_shape.top, source_shape.width, source_shape.height, str(source_shape.shape_type)) != (candidate_shape.left, candidate_shape.top, candidate_shape.width, candidate_shape.height, str(candidate_shape.shape_type)):
            return False
        if candidate_shape.text != expected[index]:
            return False
    return True


def media_hashes(path: Path) -> Counter:
    with zipfile.ZipFile(path) as archive:
        return Counter(hashlib.sha256(archive.read(name)).hexdigest() for name in archive.namelist() if name.startswith("ppt/media/") and not name.endswith("/"))


def render(path: Path, outdir: Path) -> tuple[int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "profile"
    profile.mkdir()
    completed = subprocess.run(["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(path)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300, check=False)
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return 0, completed.stdout[-500:]
    info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=30, check=False).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    return (int(match.group(1)) if match else 0), completed.stdout[-500:]


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {"criterion_id": identifier, "score": 1.0 if passed else 0.0, "evidence": evidence, "method": "deterministic", "raw": {"description": description, "weight": weight, "type": kind, "passed": bool(passed)}}


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [item["criterion_id"] for item in criteria if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0 and item["raw"]["passed"])
    penalty = sum(abs(item["raw"]["weight"]) for item in criteria if item["raw"]["weight"] < 0 and not item["raw"]["passed"])
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [sha256(path)] if path.is_file() else [], "criterion_results": criteria, "reward": reward, "passed": reward >= 0.8 and not hurdles, "hurdle_failures": hurdles}


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        return make_result(path, [result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配") for identifier in SPECS])
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        return make_result(path, [result_item(identifier, False, f"PPTX 打开失败：{exc}") for identifier in SPECS])

    criteria = []
    package_ok, package_evidence = package_checks(path)
    size_ok = len(candidate.slides) == 3 and (source.slide_width, source.slide_height) == (candidate.slide_width, candidate.slide_height)
    criteria.append(result_item("R001", package_ok and size_ok, f"页数={len(candidate.slides)}，尺寸一致={size_ok}，包={json.dumps(package_evidence, ensure_ascii=False)}"))

    texts = all_text(candidate)
    required_ok = all(texts.count(text) == 1 for text in REQUIRED_TEXT)
    deleted_ok = not any(text.casefold() in {"team b performance", "performance chart", "team a performance", "images and captions".lower()} for text in texts if text not in REQUIRED_TEXT)
    charts = [shape.shape_id for slide in candidate.slides for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.CHART]
    title_ok = required_ok and deleted_ok and not charts
    criteria.append(result_item("R002", title_ok, f"必要标题={required_ok}，删页文字残留={not deleted_ok}，图表={charts}"))

    table_shape = only_table(candidate)
    if table_shape is not None:
        facts, table_errors, table_detail = table_facts(table_shape)
        header_ok = [cell.text.strip() for cell in table_shape.table.rows[0].cells] == ["Category", "Team A", "Team B"]
        facts_ok = facts == EXPECTED_FACTS and not table_errors and header_ok
    else:
        facts, table_errors, table_detail, facts_ok = {}, ["缺少第 2 页唯一表格"], {}, False
    criteria.append(result_item("R003", facts_ok, f"事实={json.dumps({str(key): value for key, value in facts.items()}, ensure_ascii=False)}，错误={table_errors}，表={table_detail}"))

    old_list_tokens = ("Manufacturing – 20", "Logistics – 35", "Manufacturing – 30", "Sales – 20")
    slide2_text = "\n".join(shape.text for shape in candidate.slides[1].shapes if getattr(shape, "has_text_frame", False)) if len(candidate.slides) >= 2 else ""
    old_list_removed = not any(token in slide2_text for token in old_list_tokens)
    table_primary = table_shape is not None and table_shape.width >= candidate.slide_width * 0.7 and table_shape.width * table_shape.height >= candidate.slide_width * candidate.slide_height * 0.25
    editable_table_ok = table_shape is not None and old_list_removed and table_primary and not charts
    criteria.append(result_item("R004", editable_table_ok, f"旧列表移除={old_list_removed}，表格为主={table_primary}，图表={charts}"))

    image_ok, caption_ok, final_detail = final_slide_analysis(candidate)
    criteria.append(result_item("R005", image_ok, json.dumps(final_detail, ensure_ascii=False)))
    criteria.append(result_item("R006", caption_ok, json.dumps(final_detail, ensure_ascii=False)))

    combined_text = "\n".join(texts).lower()
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker in combined_text]
    media_ok = safe(lambda: media_hashes(path) == Counter(IMAGE_ORDER), False)
    cover_ok = source_cover_preserved(source, candidate)
    full_slide_pictures = [shape.shape_id for slide in candidate.slides for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.width * shape.height >= candidate.slide_width * candidate.slide_height * 0.8]
    preservation_ok = media_ok and cover_ok and not full_slide_pictures and not markers
    criteria.append(result_item("R007", preservation_ok, f"媒体={media_ok}，封面={cover_ok}，整页图={full_slide_pictures}，标记={markers}"))

    pages, render_log = 0, ""
    try:
        with tempfile.TemporaryDirectory(prefix="pptarena074-render-") as temporary:
            pages, render_log = render(path, Path(temporary))
    except Exception as exc:
        render_log = f"渲染异常：{exc}"
    delivered = sorted(str(item.relative_to(artifact_dir)) for item in artifact_dir.rglob("*") if item.is_file())
    hygiene_ok = pages == 3 and delivered == [OUTPUT_NAME] and package_ok
    criteria.append(result_item("R008", hygiene_ok, f"渲染页数={pages}，文件={delivered}，日志={render_log}"))

    severe = not (title_ok and facts_ok and editable_table_ok and image_ok and caption_ok and preservation_ok)
    criteria.append(result_item("R009", not severe, f"严重错误={severe}"))
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    path = args.artifact_dir / OUTPUT_NAME
    result = evaluate(path, args.artifact_dir) if path.is_file() else make_result(path, [result_item(identifier, False, f"缺失指定交付物：{OUTPUT_NAME}") for identifier in SPECS])
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{result['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
