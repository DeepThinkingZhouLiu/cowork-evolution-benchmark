#!/usr/bin/env python3
"""只为 PPTArena Case 82 构造可重放的合法变体、边界例和负例。"""

from __future__ import annotations

import argparse
import re
import shutil
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Inches


OUTPUT_NAME = "AddSlideBorderGroundTruthB_AccountantCVResume.pptx"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
SLIDE_WIDTH = 7_556_500
SLIDE_HEIGHT = 10_693_400
EMU_PER_INCH = 914_400


def qn(namespace: str, local: str) -> str:
    return f"{{{namespace}}}{local}"


def output_path(root: Path, case: str) -> Path:
    target = root / case / OUTPUT_NAME
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def serialize(root) -> bytes:
    return etree.tostring(
        root,
        encoding="UTF-8",
        xml_declaration=True,
        standalone=True,
    )


def rewrite_archive(source: Path, output: Path, transform) -> None:
    with zipfile.ZipFile(source) as archive_in, zipfile.ZipFile(output, "w") as archive_out:
        for info in archive_in.infolist():
            archive_out.writestr(info, transform(info.filename, archive_in.read(info.filename)))


def slide_number(name: str) -> int | None:
    match = re.fullmatch(r"ppt/slides/slide(\d+)\.xml", name)
    return int(match.group(1)) if match else None


def next_shape_id(root) -> int:
    values = [
        int(value)
        for value in root.xpath(".//p:cNvPr/@id", namespaces={"p": P_NS})
        if value.isdigit()
    ]
    return max(values, default=1) + 1


def add_shape_style(parent, *, connector: bool) -> None:
    style = etree.SubElement(parent, qn(P_NS, "style"))
    line_ref = etree.SubElement(style, qn(A_NS, "lnRef"), idx="2" if connector else "1")
    etree.SubElement(line_ref, qn(A_NS, "schemeClr"), val="accent1")
    fill_ref = etree.SubElement(style, qn(A_NS, "fillRef"), idx="0" if connector else "3")
    etree.SubElement(fill_ref, qn(A_NS, "schemeClr"), val="accent1")
    effect_ref = etree.SubElement(style, qn(A_NS, "effectRef"), idx="1" if connector else "2")
    etree.SubElement(effect_ref, qn(A_NS, "schemeClr"), val="accent1")
    font_ref = etree.SubElement(style, qn(A_NS, "fontRef"), idx="minor")
    etree.SubElement(font_ref, qn(A_NS, "schemeClr"), val="tx1" if connector else "lt1")


def add_rectangle(
    root,
    *,
    margins: tuple[float, float, float, float],
    color: str,
    width_emu: int,
    dash: str | None = None,
    fill: str | None = None,
    width_ratio: float = 1.0,
) -> int:
    left, top, right, bottom = [round(value * EMU_PER_INCH) for value in margins]
    extent_width = round((SLIDE_WIDTH - left - right) * width_ratio)
    extent_height = SLIDE_HEIGHT - top - bottom
    shape_id = next_shape_id(root)
    tree = root.xpath("./p:cSld/p:spTree", namespaces={"p": P_NS})[0]
    shape = etree.SubElement(tree, qn(P_NS, "sp"))
    non_visual = etree.SubElement(shape, qn(P_NS, "nvSpPr"))
    etree.SubElement(non_visual, qn(P_NS, "cNvPr"), id=str(shape_id), name=f"Border Rectangle {shape_id}")
    etree.SubElement(non_visual, qn(P_NS, "cNvSpPr"))
    etree.SubElement(non_visual, qn(P_NS, "nvPr"))
    properties = etree.SubElement(shape, qn(P_NS, "spPr"))
    transform = etree.SubElement(properties, qn(A_NS, "xfrm"))
    etree.SubElement(transform, qn(A_NS, "off"), x=str(left), y=str(top))
    etree.SubElement(transform, qn(A_NS, "ext"), cx=str(extent_width), cy=str(extent_height))
    geometry = etree.SubElement(properties, qn(A_NS, "prstGeom"), prst="rect")
    etree.SubElement(geometry, qn(A_NS, "avLst"))
    if fill is None:
        etree.SubElement(properties, qn(A_NS, "noFill"))
    else:
        solid_fill = etree.SubElement(properties, qn(A_NS, "solidFill"))
        etree.SubElement(solid_fill, qn(A_NS, "srgbClr"), val=fill)
    line = etree.SubElement(properties, qn(A_NS, "ln"), w=str(width_emu))
    line_fill = etree.SubElement(line, qn(A_NS, "solidFill"))
    etree.SubElement(line_fill, qn(A_NS, "srgbClr"), val=color)
    if dash is not None:
        etree.SubElement(line, qn(A_NS, "prstDash"), val=dash)
    add_shape_style(shape, connector=False)
    text_body = etree.SubElement(shape, qn(P_NS, "txBody"))
    etree.SubElement(text_body, qn(A_NS, "bodyPr"), rtlCol="0", anchor="ctr")
    etree.SubElement(text_body, qn(A_NS, "lstStyle"))
    paragraph = etree.SubElement(text_body, qn(A_NS, "p"))
    etree.SubElement(paragraph, qn(A_NS, "pPr"), algn="ctr")
    return shape_id


def add_connector(root, *, x: int, y: int, width: int, height: int, color: str, width_emu: int) -> int:
    shape_id = next_shape_id(root)
    tree = root.xpath("./p:cSld/p:spTree", namespaces={"p": P_NS})[0]
    shape = etree.SubElement(tree, qn(P_NS, "cxnSp"))
    non_visual = etree.SubElement(shape, qn(P_NS, "nvCxnSpPr"))
    etree.SubElement(non_visual, qn(P_NS, "cNvPr"), id=str(shape_id), name=f"Border Connector {shape_id}")
    etree.SubElement(non_visual, qn(P_NS, "cNvCxnSpPr"))
    etree.SubElement(non_visual, qn(P_NS, "nvPr"))
    properties = etree.SubElement(shape, qn(P_NS, "spPr"))
    transform = etree.SubElement(properties, qn(A_NS, "xfrm"))
    etree.SubElement(transform, qn(A_NS, "off"), x=str(x), y=str(y))
    etree.SubElement(transform, qn(A_NS, "ext"), cx=str(width), cy=str(height))
    geometry = etree.SubElement(properties, qn(A_NS, "prstGeom"), prst="line")
    etree.SubElement(geometry, qn(A_NS, "avLst"))
    line = etree.SubElement(properties, qn(A_NS, "ln"), w=str(width_emu))
    fill = etree.SubElement(line, qn(A_NS, "solidFill"))
    etree.SubElement(fill, qn(A_NS, "srgbClr"), val=color)
    add_shape_style(shape, connector=True)
    return shape_id


def add_four_lines(root, *, margin: float, color: str, width_emu: int, omit_right: bool = False) -> None:
    offset = round(margin * EMU_PER_INCH)
    horizontal = SLIDE_WIDTH - 2 * offset
    vertical = SLIDE_HEIGHT - 2 * offset
    add_connector(root, x=offset, y=offset, width=horizontal, height=0, color=color, width_emu=width_emu)
    add_connector(root, x=offset, y=SLIDE_HEIGHT - offset, width=horizontal, height=0, color=color, width_emu=width_emu)
    add_connector(root, x=offset, y=offset, width=0, height=vertical, color=color, width_emu=width_emu)
    if not omit_right:
        add_connector(root, x=SLIDE_WIDTH - offset, y=offset, width=0, height=vertical, color=color, width_emu=width_emu)


def rectangle_variant(source: Path, output: Path, configs: list[list[dict]]) -> None:
    def transform(name: str, data: bytes) -> bytes:
        number = slide_number(name)
        if number is None:
            return data
        root = etree.fromstring(data)
        for config in configs[number - 1]:
            add_rectangle(root, **config)
        return serialize(root)

    rewrite_archive(source, output, transform)


def four_line_variant(source: Path, output: Path, *, omit_right: bool = False) -> None:
    def transform(name: str, data: bytes) -> bytes:
        number = slide_number(name)
        if number is None:
            return data
        root = etree.fromstring(data)
        add_four_lines(root, margin=0.25, color="444444", width_emu=12_700, omit_right=omit_right)
        return serialize(root)

    rewrite_archive(source, output, transform)


def base_config(**overrides) -> dict:
    config = {
        "margins": (0.2, 0.2, 0.2, 0.2),
        "color": "555555",
        "width_emu": 9_144,
    }
    config.update(overrides)
    return config


def repacked_content_types(golden: Path, output: Path) -> None:
    rewrite_archive(
        golden,
        output,
        lambda name, data: data.replace(b"</Types>", b"  \n</Types>") if name == "[Content_Types].xml" else data,
    )


def changed_text(golden: Path, output: Path) -> None:
    def transform(name: str, data: bytes) -> bytes:
        if name != "ppt/slides/slide1.xml":
            return data
        root = etree.fromstring(data)
        nodes = root.xpath(".//a:t", namespaces={"a": A_NS})
        nodes[0].text = f"{nodes[0].text} X"
        return serialize(root)

    rewrite_archive(golden, output, transform)


def moved_original(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    presentation.slides[0].shapes[0].left += Inches(0.2)
    presentation.save(output)


def changed_image(golden: Path, output: Path) -> None:
    media = []
    with zipfile.ZipFile(golden) as archive:
        media = sorted(name for name in archive.namelist() if name.startswith("ppt/media/"))
        replacement = archive.read(media[1])
    rewrite_archive(golden, output, lambda name, data: replacement if name == media[0] else data)


def hardcoded_marker(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    box = presentation.slides[0].shapes.add_textbox(Inches(0.3), Inches(0.3), Inches(2), Inches(0.25))
    box.text = "solution/Golden"
    presentation.save(output)


def extra_slide(golden: Path, output: Path) -> None:
    presentation = Presentation(golden)
    presentation.slides.add_slide(presentation.slide_layouts[0])
    presentation.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()

    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))

    rectangle_variant(
        args.source,
        output_path(args.outdir, "valid-gray-rectangle"),
        [[base_config(margins=(0.3, 0.3, 0.3, 0.3), color="666666", width_emu=15_875)] for _ in range(4)],
    )
    four_line_variant(args.source, output_path(args.outdir, "valid-four-lines"))
    repacked_content_types(args.golden, output_path(args.outdir, "repacked-content-types"))

    rectangle_variant(args.source, output_path(args.outdir, "missing-slide4"), [[base_config()], [base_config()], [base_config()], []])
    rectangle_variant(args.source, output_path(args.outdir, "only-slide1"), [[base_config()], [], [], []])
    rectangle_variant(args.source, output_path(args.outdir, "double-slide2"), [[base_config()], [base_config(), base_config(margins=(0.3, 0.3, 0.3, 0.3))], [base_config()], [base_config()]])
    rectangle_variant(args.source, output_path(args.outdir, "filled-rectangle"), [[base_config(fill="FFFFFF")] for _ in range(4)])
    rectangle_variant(args.source, output_path(args.outdir, "too-thick"), [[base_config(width_emu=63_500)] for _ in range(4)])
    rectangle_variant(args.source, output_path(args.outdir, "too-light"), [[base_config(color="CCCCCC")] for _ in range(4)])
    rectangle_variant(args.source, output_path(args.outdir, "red-border"), [[base_config(color="880000")] for _ in range(4)])
    rectangle_variant(args.source, output_path(args.outdir, "dashed-border"), [[base_config(dash="dash")] for _ in range(4)])
    rectangle_variant(args.source, output_path(args.outdir, "edge-touching"), [[base_config(margins=(0.0, 0.0, 0.0, 0.0))] for _ in range(4)])
    rectangle_variant(
        args.source,
        output_path(args.outdir, "nonuniform-pages"),
        [[base_config(margins=(m, m, m, m))] for m in (0.2, 0.3, 0.2, 0.4)],
    )
    rectangle_variant(args.source, output_path(args.outdir, "lopsided-margin"), [[base_config(margins=(0.2, 0.2, 0.4, 0.2))] for _ in range(4)])
    rectangle_variant(args.source, output_path(args.outdir, "small-content-box"), [[base_config(width_ratio=0.5)] for _ in range(4)])
    four_line_variant(args.source, output_path(args.outdir, "missing-right-line"), omit_right=True)
    changed_text(args.golden, output_path(args.outdir, "changed-text"))
    moved_original(args.golden, output_path(args.outdir, "moved-original"))
    changed_image(args.golden, output_path(args.outdir, "changed-image"))
    hardcoded_marker(args.golden, output_path(args.outdir, "hardcoded-marker"))
    extra_slide(args.golden, output_path(args.outdir, "extra-slide"))


if __name__ == "__main__":
    main()
