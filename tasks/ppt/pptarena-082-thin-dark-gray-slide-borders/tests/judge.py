#!/usr/bin/env python3
"""PPTArena Case 82 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import statistics
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from lxml import etree
from PIL import Image, ImageChops
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE


OUTPUT_NAME = "AddSlideBorderGroundTruthB_AccountantCVResume.pptx"
INPUT_NAME = "AddSlideBorderTestB_AccountantCVResume.pptx"
INPUT_SHA256 = "bff930968572cce3f83f3bbb3eb26889090edfb7aceb681b109958ea78216ff2"
TASK_ID = "cowork-evo/pptarena-082-thin-dark-gray-slide-borders"
RUBRIC_VERSION = "pptarena-case-082-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-082-deterministic-v0.1"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
EMU_PER_INCH = 914_400
SOURCE_SHAPE_COUNTS = [58, 57, 56, 7]
SUSPICIOUS_MARKERS = [
    "solution/golden",
    "/root/",
    "judge-result",
    "artifact_dir",
    "pptarena-case-082",
]

SPECS = {
    "R001": ("指定 PPTX 可打开、页面尺寸不变且总页数为 4", 6, "hurdle"),
    "R002": ("4/4 页都有唯一闭合的原生矩形或四直线边框", 18, "hurdle"),
    "R003": ("边框四侧内缩 0.1–0.5 英寸，单页对边和跨页位置均匀", 14, "hurdle"),
    "R004": ("边框为无填充的 0.5–2pt 深灰色实线，跨页样式一致", 14, "hurdle"),
    "R005": ("除合法边框外，所有原对象、文字、几何、样式、顺序、版式和备注保持", 18, "hurdle"),
    "R006": ("页面可渲染，除边框线条区外与冻结输入高度一致", 10, "hurdle"),
    "R007": ("4 页渲染图的预期四边均有连续深灰框，且不压住原内容", 5, "hurdle"),
    "R008": ("原对象和 41 个媒体保持，只增加 4 个矩形或 16 条边线", 8, "hurdle"),
    "R009": ("输出目录干净且不含答案路径或未授权文字", 5, "positive"),
    "R010": ("移除合法边框后内容 XML 和关键包部件字节级保持", 2, "positive"),
    "R011": ("缺页边框、样式/位置越界或原内容、布局、可编辑性被破坏", -20, "penalty"),
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def qn(namespace: str, local: str) -> str:
    return f"{{{namespace}}}{local}"


def locate_source() -> Path | None:
    candidates = [
        Path("/app/input_files") / INPUT_NAME,
        Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME,
    ]
    return next((path for path in candidates if path.is_file()), None)


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def rounded_geometry(shape) -> tuple:
    return tuple(
        round(int(value) / 1000) * 1000
        for value in (shape.left, shape.top, shape.width, shape.height)
    )


def color_signature(color) -> tuple:
    try:
        color_type = int(color.type) if color.type is not None else None
    except Exception:
        color_type = str(getattr(color, "type", None))
    try:
        rgb = str(color.rgb) if color.rgb is not None else None
    except Exception:
        rgb = None
    try:
        theme = int(color.theme_color) if color.theme_color is not None else None
    except Exception:
        theme = None
    try:
        brightness = round(color.brightness, 5)
    except Exception:
        brightness = None
    return color_type, rgb, theme, brightness


def font_signature(font, run=None) -> tuple:
    try:
        language_value = font.language_id
        language = int(language_value) if language_value is not None else None
    except Exception:
        language = None
    try:
        hyperlink = run.hyperlink.address if run is not None else None
    except Exception:
        hyperlink = None
    return (
        font.name,
        round(font.size.pt, 4) if font.size else None,
        bool(font.bold),
        bool(font.italic),
        bool(font.underline),
        language,
        color_signature(font.color),
        hyperlink,
    )


def text_signature(shape) -> tuple:
    text_frame = shape.text_frame
    paragraphs = []
    for paragraph in text_frame.paragraphs:
        runs = tuple((run.text, font_signature(run.font, run)) for run in paragraph.runs)
        paragraphs.append(
            (
                paragraph.text,
                paragraph.level,
                str(paragraph.alignment),
                str(paragraph.line_spacing),
                str(paragraph.space_before),
                str(paragraph.space_after),
                runs,
            )
        )
    return (
        (
            text_frame.margin_left,
            text_frame.margin_right,
            text_frame.margin_top,
            text_frame.margin_bottom,
            str(text_frame.vertical_anchor),
            text_frame.word_wrap,
            str(text_frame.auto_size),
        ),
        tuple(paragraphs),
    )


def picture_signature(shape) -> tuple:
    try:
        blob = hashlib.sha256(shape.image.blob).hexdigest()
    except Exception as exc:
        blob = f"unreadable:{type(exc).__name__}"
    return (
        blob,
        round(float(shape.crop_left), 6),
        round(float(shape.crop_right), 6),
        round(float(shape.crop_top), 6),
        round(float(shape.crop_bottom), 6),
    )


def shape_signature(shape) -> tuple:
    text = text_signature(shape) if getattr(shape, "has_text_frame", False) else ()
    picture = picture_signature(shape) if shape.shape_type == MSO_SHAPE_TYPE.PICTURE else ()
    children = (
        tuple(shape_signature(child) for child in shape.shapes)
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP
        else ()
    )
    return (
        int(shape.shape_id),
        shape.name,
        str(shape.shape_type),
        rounded_geometry(shape),
        round(float(shape.rotation or 0), 4),
        text,
        picture,
        children,
    )


def notes_signature(presentation: Presentation) -> tuple:
    values = []
    for slide in presentation.slides:
        try:
            values.append(slide.notes_slide.notes_text_frame.text)
        except Exception:
            values.append("")
    return tuple(values)


def line_style(shape) -> dict:
    try:
        width_pt = shape.line.width.pt if shape.line.width is not None else None
    except Exception:
        width_pt = None
    try:
        rgb_text = str(shape.line.color.rgb) if shape.line.color.rgb is not None else None
    except Exception:
        rgb_text = None
    try:
        theme = str(shape.line.color.theme_color) if shape.line.color.theme_color is not None else None
    except Exception:
        theme = None
    try:
        dash = str(shape.line.dash_style) if shape.line.dash_style is not None else None
    except Exception:
        dash = None
    rgb = None
    if rgb_text and re.fullmatch(r"[0-9A-Fa-f]{6}", rgb_text):
        rgb = tuple(int(rgb_text[index : index + 2], 16) for index in (0, 2, 4))
    dark_gray = False
    if rgb is not None:
        dark_gray = (
            30 <= statistics.mean(rgb) <= 150
            and max(rgb) - min(rgb) <= 25
        )
    elif theme and any(value in theme.upper() for value in ("DARK_1", "TEXT_1")):
        dark_gray = True
    solid = dash is None or "SOLID" in dash.upper()
    return {
        "width_pt": round(width_pt, 4) if width_pt is not None else None,
        "rgb": list(rgb) if rgb is not None else None,
        "theme": theme,
        "dash": dash,
        "thin": width_pt is not None and 0.5 <= width_pt <= 2.0,
        "dark_gray": dark_gray,
        "solid": solid,
    }


def no_fill(shape) -> bool:
    try:
        fill_type = str(shape.fill.type)
    except Exception:
        return False
    return shape.fill.type is None or "BACKGROUND" in fill_type.upper()


def rectangle_candidate(shape, page_width: int, page_height: int) -> bool:
    try:
        rectangle = (
            shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
            and shape.auto_shape_type == MSO_SHAPE.RECTANGLE
        )
    except Exception:
        rectangle = False
    return (
        rectangle
        and (not getattr(shape, "has_text_frame", False) or not shape.text.strip())
        and shape.width >= page_width * 0.80
        and shape.height >= page_height * 0.80
        and 0 <= shape.left <= page_width * 0.10
        and 0 <= shape.top <= page_height * 0.10
        and shape.left + shape.width <= page_width
        and shape.top + shape.height <= page_height
    )


def edge_line_candidates(slide, page_width: int, page_height: int) -> list:
    values = []
    for shape in slide.shapes:
        if shape.shape_type != MSO_SHAPE_TYPE.LINE:
            continue
        horizontal = shape.width >= page_width * 0.80 and abs(shape.height) <= EMU_PER_INCH * 0.03
        vertical = shape.height >= page_height * 0.80 and abs(shape.width) <= EMU_PER_INCH * 0.03
        if not (horizontal or vertical):
            continue
        if horizontal:
            y = shape.top
            near_edge = 0.08 * EMU_PER_INCH <= y <= 0.55 * EMU_PER_INCH or (
                page_height - 0.55 * EMU_PER_INCH <= y <= page_height - 0.08 * EMU_PER_INCH
            )
        else:
            x = shape.left
            near_edge = 0.08 * EMU_PER_INCH <= x <= 0.55 * EMU_PER_INCH or (
                page_width - 0.55 * EMU_PER_INCH <= x <= page_width - 0.08 * EMU_PER_INCH
            )
        if near_edge:
            values.append(shape)
    return values


def inventory_frame(slide, page_width: int, page_height: int) -> dict:
    rectangles = [
        shape
        for shape in slide.shapes
        if rectangle_candidate(shape, page_width, page_height)
    ]
    lines = edge_line_candidates(slide, page_width, page_height)
    frame = {
        "mode": None,
        "shape_ids": [],
        "margins_in": None,
        "styles": [],
        "structural_passed": False,
        "geometry_passed": False,
        "style_passed": False,
        "no_fill_passed": False,
    }
    if len(rectangles) == 1 and not lines:
        shape = rectangles[0]
        left = shape.left
        top = shape.top
        right = page_width - (shape.left + shape.width)
        bottom = page_height - (shape.top + shape.height)
        frame.update(
            {
                "mode": "rectangle",
                "shape_ids": [int(shape.shape_id)],
                "margins_in": [value / EMU_PER_INCH for value in (left, top, right, bottom)],
                "styles": [line_style(shape)],
                "structural_passed": True,
                "no_fill_passed": no_fill(shape),
            }
        )
    elif not rectangles and len(lines) == 4:
        horizontal = sorted(
            [shape for shape in lines if shape.width >= page_width * 0.80],
            key=lambda shape: shape.top,
        )
        vertical = sorted(
            [shape for shape in lines if shape.height >= page_height * 0.80],
            key=lambda shape: shape.left,
        )
        if len(horizontal) == 2 and len(vertical) == 2:
            left = vertical[0].left
            right = page_width - vertical[1].left
            top = horizontal[0].top
            bottom = page_height - horizontal[1].top
            endpoints_ok = (
                abs(horizontal[0].left - left) <= EMU_PER_INCH * 0.05
                and abs(horizontal[1].left - left) <= EMU_PER_INCH * 0.05
                and abs(horizontal[0].left + horizontal[0].width - (page_width - right)) <= EMU_PER_INCH * 0.05
                and abs(horizontal[1].left + horizontal[1].width - (page_width - right)) <= EMU_PER_INCH * 0.05
                and abs(vertical[0].top - top) <= EMU_PER_INCH * 0.05
                and abs(vertical[1].top - top) <= EMU_PER_INCH * 0.05
                and abs(vertical[0].top + vertical[0].height - (page_height - bottom)) <= EMU_PER_INCH * 0.05
                and abs(vertical[1].top + vertical[1].height - (page_height - bottom)) <= EMU_PER_INCH * 0.05
            )
            frame.update(
                {
                    "mode": "four-lines",
                    "shape_ids": [int(shape.shape_id) for shape in lines],
                    "margins_in": [value / EMU_PER_INCH for value in (left, top, right, bottom)],
                    "styles": [line_style(shape) for shape in lines],
                    "structural_passed": endpoints_ok,
                    "no_fill_passed": True,
                }
            )
    if frame["margins_in"] is not None:
        margins = frame["margins_in"]
        frame["geometry_passed"] = (
            all(0.1 <= value <= 0.5 for value in margins)
            and abs(margins[0] - margins[2]) <= 0.05
            and abs(margins[1] - margins[3]) <= 0.05
        )
    styles = frame["styles"]
    frame["style_passed"] = bool(styles) and all(
        item["thin"] and item["dark_gray"] and item["solid"]
        for item in styles
    )
    return frame


def cross_slide_consistency(frames: list[dict]) -> tuple[bool, dict]:
    if len(frames) != 4 or any(frame["margins_in"] is None for frame in frames):
        return False, {"reason": "边框几何不完整"}
    margin_columns = list(zip(*(frame["margins_in"] for frame in frames)))
    margin_spreads = [max(values) - min(values) for values in margin_columns]
    widths = [item["width_pt"] for frame in frames for item in frame["styles"]]
    width_spread = max(widths) - min(widths) if widths and all(value is not None for value in widths) else float("inf")
    colors = [
        tuple(item["rgb"]) if item["rgb"] is not None else (item["theme"],)
        for frame in frames
        for item in frame["styles"]
    ]
    colors_consistent = len(set(colors)) == 1
    passed = max(margin_spreads) <= 0.03 and width_spread <= 0.15 and colors_consistent
    return passed, {
        "margin_spreads_in": [round(value, 5) for value in margin_spreads],
        "width_spread_pt": round(width_spread, 5) if math.isfinite(width_spread) else None,
        "colors": colors,
    }


def content_inside_frame(source: Presentation, frames: list[dict]) -> tuple[bool, list]:
    details = []
    passed = True
    for index, (slide, frame) in enumerate(zip(source.slides, frames), start=1):
        if frame["margins_in"] is None:
            details.append({"slide": index, "reason": "无边框几何"})
            passed = False
            continue
        left, top, right, bottom = [value * EMU_PER_INCH for value in frame["margins_in"]]
        crossing = []
        for shape in slide.shapes:
            tolerance = EMU_PER_INCH * 0.02
            if (
                shape.left < left - tolerance
                or shape.top < top - tolerance
                or shape.left + shape.width > source.slide_width - right + tolerance
                or shape.top + shape.height > source.slide_height - bottom + tolerance
            ):
                crossing.append(int(shape.shape_id))
        details.append({"slide": index, "crossing_original_shape_ids": crossing})
        passed = passed and not crossing
    return passed, details


def package_ok(path: Path) -> tuple[bool, str]:
    try:
        with zipfile.ZipFile(path) as archive:
            broken = archive.testzip()
            names = set(archive.namelist())
            required = {
                "[Content_Types].xml",
                "ppt/presentation.xml",
                "ppt/slides/slide1.xml",
                "ppt/slides/slide2.xml",
                "ppt/slides/slide3.xml",
                "ppt/slides/slide4.xml",
            }
            missing = sorted(required - names)
            macros = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
        return broken is None and not missing and not macros, (
            f"损坏成员={broken}，缺失={missing}，宏={macros}"
        )
    except Exception as exc:
        return False, f"ZIP 检查失败：{exc}"


def media_hashes(path: Path) -> Counter[str]:
    with zipfile.ZipFile(path) as archive:
        return Counter(
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("ppt/media/")
        )


def member_names(path: Path) -> set[str]:
    with zipfile.ZipFile(path) as archive:
        return {name for name in archive.namelist() if not name.endswith("/")}


def critical_hashes(path: Path) -> dict[str, str]:
    excluded = {f"ppt/slides/slide{index}.xml" for index in range(1, 5)}
    with zipfile.ZipFile(path) as archive:
        return {
            name: hashlib.sha256(archive.read(name)).hexdigest()
            for name in sorted(archive.namelist())
            if not name.endswith("/") and name not in excluded and not name.startswith("docProps/")
        }


def normalized_slides_hash(path: Path, frame_ids: list[list[int]]) -> str:
    values = []
    with zipfile.ZipFile(path) as archive:
        for slide_index, ids in enumerate(frame_ids, start=1):
            root = etree.fromstring(archive.read(f"ppt/slides/slide{slide_index}.xml"))
            for shape_id in ids:
                matches = root.xpath(
                    f"./p:cSld/p:spTree/*[p:nvSpPr/p:cNvPr/@id='{shape_id}' or p:nvCxnSpPr/p:cNvPr/@id='{shape_id}' or p:nvPicPr/p:cNvPr/@id='{shape_id}' or p:nvGrpSpPr/p:cNvPr/@id='{shape_id}']",
                    namespaces={"p": P_NS},
                )
                for shape in matches:
                    shape.getparent().remove(shape)
            c_sld = root.find(qn(P_NS, "cSld"))
            values.append(etree.tostring(c_sld, method="c14n"))
    return hashlib.sha256(b"\n".join(values)).hexdigest()


def render_presentation(path: Path, outdir: Path) -> tuple[list[Path], int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "lo-profile"
    profile.mkdir()
    completed = subprocess.run(
        [
            "soffice",
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(outdir),
            str(path),
        ],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return [], 0, f"转 PDF 失败：{completed.stdout[-500:]}"
    info = subprocess.run(
        ["pdfinfo", str(pdf)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=30,
        check=False,
    ).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    pages = int(match.group(1)) if match else 0
    prefix = outdir / "slide"
    raster = subprocess.run(
        ["pdftoppm", "-jpeg", "-r", "120", str(pdf), str(prefix)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=240,
        check=False,
    )
    images = sorted(outdir.glob("slide-*.jpg")) if raster.returncode == 0 else []
    return images, pages, (
        f"页数={pages}，渲染图={len(images)}，"
        f"soffice={completed.returncode}，pdftoppm={raster.returncode}"
    )


def frame_pixel_positions(image: Image.Image, margins_in: list[float], slide_size: tuple[int, int]) -> tuple[int, int, int, int]:
    left = round(margins_in[0] * EMU_PER_INCH * image.width / slide_size[0])
    top = round(margins_in[1] * EMU_PER_INCH * image.height / slide_size[1])
    right = image.width - round(margins_in[2] * EMU_PER_INCH * image.width / slide_size[0])
    bottom = image.height - round(margins_in[3] * EMU_PER_INCH * image.height / slide_size[1])
    return left, top, right, bottom


def outside_frame_similarity(
    source_image: Image.Image,
    candidate_image: Image.Image,
    margins_in: list[float],
    slide_size: tuple[int, int],
) -> tuple[float, dict]:
    source = source_image.convert("RGB")
    candidate = candidate_image.convert("RGB")
    if source.size != candidate.size:
        return 0.0, {"reason": "渲染尺寸不一致"}
    left, top, right, bottom = frame_pixel_positions(source, margins_in, slide_size)
    padding = max(4, round(min(source.size) * 0.004))
    mask = Image.new("L", source.size, 255)
    mask.paste(0, (left - padding, top - padding, right + padding, top + padding + 1))
    mask.paste(0, (left - padding, bottom - padding, right + padding, bottom + padding + 1))
    mask.paste(0, (left - padding, top - padding, left + padding + 1, bottom + padding))
    mask.paste(0, (right - padding, top - padding, right + padding + 1, bottom + padding))
    difference = ImageChops.difference(source, candidate)
    pixels = mask.histogram()[255]
    histogram = difference.histogram(mask)
    squared = sum(count * ((index % 256) ** 2) for index, count in enumerate(histogram))
    rms = math.sqrt(squared / (pixels * 3)) if pixels else 255.0
    return max(0.0, 1.0 - rms / 255.0), {
        "frame_pixels": [left, top, right, bottom],
        "padding": padding,
        "evaluated_pixels": pixels,
        "rms": round(rms, 6),
    }


def dark_line_coverage(image: Image.Image, margins_in: list[float], slide_size: tuple[int, int]) -> tuple[bool, dict]:
    rgb = image.convert("RGB")
    left, top, right, bottom = frame_pixel_positions(rgb, margins_in, slide_size)
    band = max(2, round(min(rgb.size) * 0.0025))

    def horizontal(y: int) -> float:
        values = []
        for x in range(left + 4, right - 4):
            pixels = [rgb.getpixel((x, yy)) for yy in range(max(0, y - band), min(rgb.height, y + band + 1))]
            values.append(min(statistics.mean(pixel) for pixel in pixels) <= 155)
        return statistics.mean(values) if values else 0.0

    def vertical(x: int) -> float:
        values = []
        for y in range(top + 4, bottom - 4):
            pixels = [rgb.getpixel((xx, y)) for xx in range(max(0, x - band), min(rgb.width, x + band + 1))]
            values.append(min(statistics.mean(pixel) for pixel in pixels) <= 155)
        return statistics.mean(values) if values else 0.0

    coverage = {
        "top": horizontal(top),
        "bottom": horizontal(bottom),
        "left": vertical(left),
        "right": vertical(right),
    }
    passed = all(value >= 0.85 for value in coverage.values())
    return passed, {key: round(value, 4) for key, value in coverage.items()}


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {
        "criterion_id": identifier,
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
        "raw": {
            "description": description,
            "weight": weight,
            "type": kind,
            "passed": bool(passed),
        },
    }


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [
        item["criterion_id"]
        for item in criteria
        if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]
    ]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(
        item["raw"]["weight"]
        for item in criteria
        if item["raw"]["weight"] > 0 and item["raw"]["passed"]
    )
    penalty = sum(
        abs(item["raw"]["weight"])
        for item in criteria
        if item["raw"]["weight"] < 0 and not item["raw"]["passed"]
    )
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [sha256(path)] if path.is_file() else [],
        "criterion_results": criteria,
        "reward": reward,
        "passed": reward >= 0.8 and not hurdles,
        "hurdle_failures": hurdles,
    }


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        return make_result(
            path,
            [result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配") for identifier in SPECS],
        )
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        return make_result(
            path,
            [result_item(identifier, False, f"PPTX 打开失败：{exc}") for identifier in SPECS],
        )

    criteria: list[dict] = []
    package_passed, package_evidence = package_ok(path)
    size_passed = (
        len(source.slides) == 4
        and len(candidate.slides) == 4
        and candidate.slide_width == source.slide_width
        and candidate.slide_height == source.slide_height
    )
    criteria.append(
        result_item(
            "R001",
            package_passed and size_passed,
            f"{package_evidence}，输入页数={len(source.slides)}，"
            f"候选页数={len(candidate.slides)}，页面="
            f"{candidate.slide_width / EMU_PER_INCH:.4f}×"
            f"{candidate.slide_height / EMU_PER_INCH:.4f}英寸",
        )
    )

    frames = [
        inventory_frame(slide, candidate.slide_width, candidate.slide_height)
        for slide in candidate.slides
    ] if len(candidate.slides) == 4 else []
    structural_passed = len(frames) == 4 and all(frame["structural_passed"] for frame in frames)
    criteria.append(
        result_item(
            "R002",
            structural_passed,
            f"边框结构={json.dumps(frames, ensure_ascii=False)}",
        )
    )

    consistency_passed, consistency = cross_slide_consistency(frames)
    geometry_passed = (
        len(frames) == 4
        and all(frame["geometry_passed"] for frame in frames)
        and consistency_passed
    )
    criteria.append(
        result_item(
            "R003",
            geometry_passed,
            f"跨页几何={json.dumps(consistency, ensure_ascii=False)}",
        )
    )

    style_passed = (
        len(frames) == 4
        and all(frame["style_passed"] and frame["no_fill_passed"] for frame in frames)
        and consistency_passed
    )
    criteria.append(
        result_item(
            "R004",
            style_passed,
            f"边框样式={json.dumps([frame.get('styles') for frame in frames], ensure_ascii=False)}，"
            f"无填充={[frame.get('no_fill_passed') for frame in frames]}",
        )
    )

    preservation_passed = False
    preservation_details = []
    if len(candidate.slides) == 4 and len(frames) == 4:
        per_slide = []
        for index, (source_slide, candidate_slide, frame) in enumerate(
            zip(source.slides, candidate.slides, frames), start=1
        ):
            ids = set(frame["shape_ids"])
            source_signatures = [shape_signature(shape) for shape in source_slide.shapes]
            candidate_signatures = [
                shape_signature(shape)
                for shape in candidate_slide.shapes
                if int(shape.shape_id) not in ids
            ]
            equal = source_signatures == candidate_signatures
            per_slide.append(equal)
            preservation_details.append(
                {
                    "slide": index,
                    "source_shapes": len(source_signatures),
                    "candidate_non_border_shapes": len(candidate_signatures),
                    "equal": equal,
                }
            )
        layouts_equal = all(
            str(source_slide.slide_layout.part.partname) == str(candidate_slide.slide_layout.part.partname)
            for source_slide, candidate_slide in zip(source.slides, candidate.slides)
        )
        notes_equal = notes_signature(source) == notes_signature(candidate)
        preservation_passed = all(per_slide) and layouts_equal and notes_equal
        preservation_details.append({"layouts_equal": layouts_equal, "notes_equal": notes_equal})
    criteria.append(
        result_item(
            "R005",
            preservation_passed,
            json.dumps(preservation_details, ensure_ascii=False),
        )
    )

    try:
        source_media = media_hashes(source_path)
        candidate_media = media_hashes(path)
        source_members = member_names(source_path)
        candidate_members = member_names(path)
    except Exception:
        source_media = candidate_media = Counter()
        source_members = candidate_members = set()
    added_counts = [len(frame["shape_ids"]) for frame in frames]
    valid_addition_pattern = added_counts == [1, 1, 1, 1] or added_counts == [4, 4, 4, 4]
    editable_passed = (
        package_passed
        and structural_passed
        and valid_addition_pattern
        and source_media == candidate_media
        and sum(candidate_media.values()) == 41
        and source_members == candidate_members
        and [len(slide.shapes) for slide in candidate.slides]
        == [count + added for count, added in zip(SOURCE_SHAPE_COUNTS, added_counts)]
    )
    criteria.append(
        result_item(
            "R008",
            editable_passed,
            f"每页新增={added_counts}，对象数={[len(slide.shapes) for slide in candidate.slides]}，"
            f"媒体哈希相等={source_media == candidate_media}，媒体数="
            f"{sum(source_media.values())}→{sum(candidate_media.values())}，包成员集相等={source_members == candidate_members}",
        )
    )

    artifact_files = sorted(
        str(item.relative_to(artifact_dir))
        for item in artifact_dir.rglob("*")
        if item.is_file()
    )
    all_text = "\n".join(
        shape.text
        for slide in candidate.slides
        for shape in iter_shapes(slide.shapes)
        if getattr(shape, "has_text_frame", False)
    )
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker.lower() in all_text.lower()]
    hygiene_passed = artifact_files == [OUTPUT_NAME] and not markers
    criteria.append(
        result_item(
            "R009",
            hygiene_passed,
            f"输出文件={artifact_files}，可疑标记={markers}",
        )
    )

    try:
        candidate_ids = [frame["shape_ids"] for frame in frames]
        source_ids = [[] for _ in range(4)]
        exact_slides = normalized_slides_hash(path, candidate_ids) == normalized_slides_hash(source_path, source_ids)
        critical_equal = critical_hashes(path) == critical_hashes(source_path)
    except Exception:
        exact_slides = critical_equal = False
    criteria.append(
        result_item(
            "R010",
            exact_slides and critical_equal,
            f"移除边框后 4 页 cSld 保持={exact_slides}，关键包部件保持={critical_equal}",
        )
    )

    inside_passed, inside_details = content_inside_frame(source, frames) if frames else (False, [])
    pre_render_ready = (
        package_passed
        and size_passed
        and structural_passed
        and geometry_passed
        and style_passed
        and preservation_passed
        and editable_passed
        and inside_passed
    )
    render_preservation_passed = False
    frame_visible_passed = False
    similarities = []
    frame_coverages = []
    render_note = "前置检查失败，未执行渲染"
    if pre_render_ready:
        with tempfile.TemporaryDirectory(prefix="pptarena082-judge-") as temporary:
            root = Path(temporary)
            source_images, source_pages, source_note = render_presentation(source_path, root / "source")
            candidate_images, candidate_pages, candidate_note = render_presentation(path, root / "candidate")
            render_note = f"源[{source_note}]；候选[{candidate_note}]"
            if len(source_images) == 4 and len(candidate_images) == 4:
                for source_image_path, candidate_image_path, frame in zip(source_images, candidate_images, frames):
                    with Image.open(source_image_path) as source_image, Image.open(candidate_image_path) as candidate_image:
                        similarity, evidence = outside_frame_similarity(
                            source_image,
                            candidate_image,
                            frame["margins_in"],
                            (source.slide_width, source.slide_height),
                        )
                        visible, coverage = dark_line_coverage(
                            candidate_image,
                            frame["margins_in"],
                            (candidate.slide_width, candidate.slide_height),
                        )
                        similarities.append({"value": similarity, **evidence})
                        frame_coverages.append({"passed": visible, **coverage})
            render_preservation_passed = (
                source_pages == 4
                and candidate_pages == 4
                and len(similarities) == 4
                and min(item["value"] for item in similarities) >= 0.9985
            )
            frame_visible_passed = (
                len(frame_coverages) == 4
                and all(item["passed"] for item in frame_coverages)
            )
    criteria.append(
        result_item(
            "R006",
            render_preservation_passed,
            f"边框区外相似度={json.dumps(similarities, ensure_ascii=False)}；{render_note}",
        )
    )
    criteria.append(
        result_item(
            "R007",
            frame_visible_passed and inside_passed,
            f"渲染边框连续度={json.dumps(frame_coverages, ensure_ascii=False)}，"
            f"原对象越过边框={json.dumps(inside_details, ensure_ascii=False)}",
        )
    )

    severe = (
        not pre_render_ready
        or not render_preservation_passed
        or not frame_visible_passed
        or bool(markers)
    )
    criteria.append(result_item("R011", not severe, f"严重错误={severe}"))
    order = {identifier: index for index, identifier in enumerate(SPECS)}
    criteria.sort(key=lambda item: order[item["criterion_id"]])
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    output = args.artifact_dir / OUTPUT_NAME
    if output.is_file():
        evaluation = evaluate(output, args.artifact_dir)
    else:
        evaluation = {
            "task_id": TASK_ID,
            "rubric_version": RUBRIC_VERSION,
            "judge_version": JUDGE_VERSION,
            "artifact_sha256": [],
            "criterion_results": [],
            "reward": 0.0,
            "passed": False,
            "hurdle_failures": ["missing_output"],
        }
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(
        json.dumps(evaluation, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    args.reward_file.write_text(f"{evaluation['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
