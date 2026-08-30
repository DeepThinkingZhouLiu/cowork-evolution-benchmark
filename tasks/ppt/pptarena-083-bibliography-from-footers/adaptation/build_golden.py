#!/usr/bin/env python3
"""PPTArena Case 83 逐页审计后的 Golden 回放器。"""

from __future__ import annotations

import argparse
import hashlib
import shutil
import zipfile
from collections import Counter
from pathlib import Path

from pptx import Presentation


INPUT_SHA256 = "17e05107cdcf497fe70bc857d7285515b1775aed43a63b32ec13458796b5670d"
UPSTREAM_SHA256 = "d3e0098d68b0bd19b513e8de2100969f573a5ce4bdf5803924de7ca27ca22136"
EXPECTED_PREFIXES = ("Agarwal", "M. Ali", "Garg", "Office", "Patrick", "Price-Haywood", "Richardson", "Vahidy", "Wosik")


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def media_hashes(path: Path) -> Counter:
    with zipfile.ZipFile(path) as archive:
        return Counter(
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("ppt/media/") and not name.endswith("/")
        )


def build(source: Path, upstream: Path, output: Path) -> None:
    if sha256(source) != INPUT_SHA256 or sha256(upstream) != UPSTREAM_SHA256:
        raise ValueError("冻结输入或上游 Ground Truth SHA256 不匹配")
    original = Presentation(source)
    answer = Presentation(upstream)
    if len(original.slides) != 23 or len(answer.slides) != 24:
        raise ValueError("页数与已审计契约不匹配")
    bibliography = answer.slides[-2]
    texts = [
        paragraph.text.strip()
        for shape in bibliography.shapes
        if getattr(shape, "has_text_frame", False)
        for paragraph in shape.text_frame.paragraphs
        if paragraph.text.strip() and paragraph.text.strip() != "Bibliography" and not paragraph.text.strip().isdigit()
    ]
    prefixes = tuple("M. Ali" if text.startswith("M. Ali") else text.split()[0] for text in texts)
    if len(texts) != 9 or prefixes != EXPECTED_PREFIXES:
        raise ValueError("上游参考文献条目不再符合 9 条排序契约")
    if media_hashes(source) != media_hashes(upstream) or len(media_hashes(source)) != 15:
        raise ValueError("上游答案未完整保留 15 个媒体资源")
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(upstream, output)
    if sha256(output) != UPSTREAM_SHA256:
        raise ValueError("Golden 回放后 SHA256 不匹配")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--upstream", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.source, args.upstream, args.output)
    print(f"{args.output}\nsha256={sha256(args.output)}")


if __name__ == "__main__":
    main()
