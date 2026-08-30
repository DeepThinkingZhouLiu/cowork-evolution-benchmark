#!/usr/bin/env python3
"""PPTArena Case 83 专用 Judge 校准样例。"""

from __future__ import annotations

import argparse
import shutil
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches


OUTPUT_NAME = "AddBibliography_GroundTruthA.pptx"


def output_path(root: Path, case: str) -> Path:
    path = root / case / OUTPUT_NAME
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def save_mutation(golden: Path, output: Path, callback) -> None:
    presentation = Presentation(golden)
    callback(presentation)
    presentation.save(output)


def bibliography_body(presentation: Presentation):
    matches = [
        shape
        for shape in presentation.slides[-2].shapes
        if getattr(shape, "has_text_frame", False)
        and sum(bool(paragraph.text.strip()) for paragraph in shape.text_frame.paragraphs) == 9
    ]
    if len(matches) != 1:
        raise ValueError("缺少唯一 Bibliography 内容框")
    return matches[0]


def repackage(golden: Path, output: Path) -> None:
    with zipfile.ZipFile(golden) as reader, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_STORED) as writer:
        for info in reader.infolist():
            writer.writestr(info.filename, reader.read(info.filename))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--upstream", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()
    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))
    shutil.copy2(args.upstream, output_path(args.outdir, "upstream-valid"))
    repackage(args.golden, output_path(args.outdir, "valid-repackage"))

    def missing(prs):
        paragraph = bibliography_body(prs).text_frame.paragraphs[2]
        paragraph._p.getparent().remove(paragraph._p)
    save_mutation(args.golden, output_path(args.outdir, "missing-citation"), missing)

    def duplicate(prs):
        paragraphs = bibliography_body(prs).text_frame.paragraphs
        paragraphs[8].text = paragraphs[7].text
    save_mutation(args.golden, output_path(args.outdir, "duplicate-citation"), duplicate)

    def wrong_order(prs):
        paragraphs = bibliography_body(prs).text_frame.paragraphs
        paragraphs[0]._p.addprevious(paragraphs[1]._p)
    save_mutation(args.golden, output_path(args.outdir, "wrong-order"), wrong_order)

    def lost_italic(prs):
        first = bibliography_body(prs).text_frame.paragraphs[0]
        for run in first.runs:
            if "N Engl J Med" in run.text:
                run.font.italic = False
    save_mutation(args.golden, output_path(args.outdir, "lost-italic"), lost_italic)

    def wrong_position(prs):
        slide_ids = prs.slides._sldIdLst
        bibliography = slide_ids[22]
        slide_ids.remove(bibliography)
        slide_ids.insert(21, bibliography)
    save_mutation(args.golden, output_path(args.outdir, "wrong-position"), wrong_position)

    def modified_footer(prs):
        shape = next(shape for shape in prs.slides[3].shapes if getattr(shape, "has_text_frame", False) and "M. Ali" in shape.text)
        shape.text = shape.text.replace("2020", "2021")
    save_mutation(args.golden, output_path(args.outdir, "modified-original-footer"), modified_footer)

    def removed_footer(prs):
        shape = next(shape for shape in prs.slides[4].shapes if getattr(shape, "has_text_frame", False) and "Wosik J" in shape.text)
        shape._element.getparent().remove(shape._element)
    save_mutation(args.golden, output_path(args.outdir, "removed-original-footer"), removed_footer)

    def changed_original(prs):
        box = prs.slides[0].shapes.add_textbox(Inches(0.1), Inches(0.1), Inches(2.0), Inches(0.2))
        box.text = "unexpected original-slide change"
    save_mutation(args.golden, output_path(args.outdir, "changed-original-slide"), changed_original)

    def media_removed(prs):
        picture = next(shape for slide in prs.slides for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
        picture._element.getparent().remove(picture._element)
    save_mutation(args.golden, output_path(args.outdir, "media-removed"), media_removed)

    placeholder = args.outdir / "flatten-placeholder.png"
    Image.new("RGB", (1200, 900), "white").save(placeholder)
    def flattened(prs):
        slide = prs.slides[-2]
        for shape in list(slide.shapes):
            shape._element.getparent().remove(shape._element)
        slide.shapes.add_picture(str(placeholder), 0, 0, width=prs.slide_width, height=prs.slide_height)
    save_mutation(args.golden, output_path(args.outdir, "flattened-bibliography"), flattened)

    def hardcoded(prs):
        box = prs.slides[-2].shapes.add_textbox(Inches(0.1), Inches(0.1), Inches(2.0), Inches(0.2))
        box.text = "solution/Golden"
    save_mutation(args.golden, output_path(args.outdir, "hardcoded-marker"), hardcoded)

    extra = args.outdir / "extra-output"
    extra.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.golden, extra / OUTPUT_NAME)
    (extra / "debug.txt").write_text("debug\n", encoding="utf-8")


if __name__ == "__main__":
    main()
