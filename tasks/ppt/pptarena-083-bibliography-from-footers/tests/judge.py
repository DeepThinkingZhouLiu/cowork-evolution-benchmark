#!/usr/bin/env python3
"""PPTArena Case 83 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import posixpath
import re
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


OUTPUT_NAME = "AddBibliography_GroundTruthA.pptx"
INPUT_NAME = "AddBibliography_TestA.pptx"
INPUT_SHA256 = "17e05107cdcf497fe70bc857d7285515b1775aed43a63b32ec13458796b5670d"
TASK_ID = "cowork-evo/pptarena-083-bibliography-from-footers"
RUBRIC_VERSION = "pptarena-case-083-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-083-deterministic-v0.1"


EXPECTED_CITATIONS = (
    "Agarwal SD, Sommers BD. Insurance Coverage after Job Loss - The Importance of the ACA during the Covid-Associated Recession [published online ahead of print, 2020 Aug 19]. N Engl J Med. 2020;10.1056/NEJMp2023312. doi:10.1056/NEJMp2023312",
    "M. Ali and C. Mangione. Health Affairs, In Press, 2020",
    "Garg S, et al. Hospitalization Rates and Characteristics of Patients Hospitalized with Laboratory-Confirmed Coronavirus Disease 2019 - COVID-NET, 14 States, March 1-30, 2020. MMWR Morb Mortal Wkly Rep. 2020;69(15):458-464. Published 2020 Apr 17. doi:10.15585/mmwr.mm6915e3",
    "Office of Health Informatics and Analytics, UCLA Health Information Technology",
    "Patrick SW, Henkhaus LE, Zickafoose JS, et al. Wellbeing of Parents and Children During the COVID-19 Pandemic: A National Survey. Pediatrics. 2020;146(4):e2020016824",
    "Price-Haywood EG et al. N Engl J Med 2020;382:2534-2543",
    "Richardson S, et al. Presenting Characteristics, Comorbidities, and Outcomes Among 5700 Patients Hospitalized With COVID-19 in the New York City Area. JAMA. 2020;323(20):2052–2059. doi:10.1001/jama.2020.6775",
    "Vahidy FS, et al. BMJ Open 2020;10:e039849. doi:10.1136/ bmjopen-2020-039849",
    "Wosik J, et al. Telehealth transformation: COVID-19 and the rise of virtual care. J Am Med Inform Assoc. 2020;27(6):957-962. doi:10.1093/jamia/ocaa067",
)

SOURCE_FOOTERS = (
    (4, "From: M. Ali and C. Mangione. Health Affairs, In Press, 2020"),
    (5, EXPECTED_CITATIONS[8]),
    (6, EXPECTED_CITATIONS[3]),
    (7, EXPECTED_CITATIONS[3]),
    (11, EXPECTED_CITATIONS[6]),
    (12, EXPECTED_CITATIONS[7]),
    (13, EXPECTED_CITATIONS[7]),
    (14, EXPECTED_CITATIONS[2]),
    (15, EXPECTED_CITATIONS[5]),
    (16, EXPECTED_CITATIONS[5]),
    (18, EXPECTED_CITATIONS[0]),
    (19, EXPECTED_CITATIONS[4]),
    (21, EXPECTED_CITATIONS[4]),
)

ITALIC_REQUIREMENTS = {
    EXPECTED_CITATIONS[0]: "N Engl J Med",
    EXPECTED_CITATIONS[1]: "In Press,",
    EXPECTED_CITATIONS[2]: "MMWR Morb Mortal Wkly Rep",
    EXPECTED_CITATIONS[6]: "JAMA.",
    EXPECTED_CITATIONS[8]: "J Am Med Inform Assoc",
}

SPECS = {
    "R001": ("PPTX 完整、保留画布且最终 24 页", 6, "hurdle"),
    "R002": ("Bibliography 倒数第 2 且最后为原 Thank you", 10, "hurdle"),
    "R003": ("9 条唯一参考文献文字准确完整", 24, "hurdle"),
    "R004": ("9 条独立段落且作者/机构字母排序", 10, "hurdle"),
    "R005": ("5 处重要斜体完整保留", 12, "hurdle"),
    "R006": ("13 处原页脚引用文字、几何与样式未改", 14, "hurdle"),
    "R007": ("原 23 页对象、备注和 15 个媒体保留", 14, "hurdle"),
    "R008": ("Bibliography 为原生可编辑文本且无答案标记", 6, "hurdle"),
    "R009": ("24 页可渲染且输出目录干净", 4, "positive"),
    "R010": ("无漏项、重复、错序、丢斜体、改页脚、错插页或压平", -20, "penalty"),
}

SUSPICIOUS_MARKERS = ["solution/golden", "judge-result", "artifact_dir", "pptarena-case-083", "/root/"]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("\u00a0", " ").replace("\v", " ")).strip()


def locate_source() -> Path | None:
    candidates = [
        Path("/app/input_files") / INPUT_NAME,
        Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME,
    ]
    return next((path for path in candidates if path.is_file()), None)


def package_checks(path: Path) -> tuple[bool, dict]:
    evidence = {"broken_member": None, "missing_targets": [], "macros": []}
    try:
        with zipfile.ZipFile(path) as archive:
            evidence["broken_member"] = archive.testzip()
            names = set(archive.namelist())
            evidence["macros"] = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
            for rels_name in sorted(name for name in names if name.endswith(".rels")):
                root = etree.fromstring(archive.read(rels_name))
                if rels_name == "_rels/.rels":
                    owner = ""
                else:
                    prefix, filename = rels_name.split("/_rels/", 1)
                    owner = posixpath.join(prefix, filename[:-5])
                base = posixpath.dirname(owner)
                for relation in root:
                    if relation.attrib.get("TargetMode") == "External":
                        continue
                    target = relation.attrib.get("Target", "")
                    resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                    if resolved not in names:
                        evidence["missing_targets"].append(f"{rels_name} -> {resolved}")
        return evidence["broken_member"] is None and not evidence["missing_targets"] and not evidence["macros"], evidence
    except Exception as exc:
        evidence["exception"] = str(exc)
        return False, evidence


def media_hashes(path: Path) -> Counter:
    with zipfile.ZipFile(path) as archive:
        return Counter(
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("ppt/media/") and not name.endswith("/")
        )


def bibliography_records(presentation: Presentation) -> list[tuple[str, object]]:
    if len(presentation.slides) < 2:
        return []
    records = []
    for shape in presentation.slides[-2].shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        for paragraph in shape.text_frame.paragraphs:
            value = normalize(paragraph.text)
            if not value or value == "Bibliography" or value.isdigit():
                continue
            records.append((value, paragraph))
    return records


def italic_substring(paragraph, substring: str) -> bool:
    normalized_runs = []
    cursor = 0
    for run in paragraph.runs:
        value = run.text.replace("\u00a0", " ").replace("\v", " ")
        normalized_runs.append((cursor, cursor + len(value), run.font.italic is True))
        cursor += len(value)
    full = "".join(run.text.replace("\u00a0", " ").replace("\v", " ") for run in paragraph.runs)
    start = full.find(substring)
    if start < 0:
        return False
    end = start + len(substring)
    covered = set()
    for left, right, italic in normalized_runs:
        if not italic:
            continue
        covered.update(range(max(start, left), min(end, right)))
    return all(index in covered for index in range(start, end) if not full[index].isspace())


def run_signature(run) -> tuple:
    font = run.font
    return (
        run.text,
        font.name,
        round(font.size.pt, 4) if font.size else None,
        font.bold,
        font.italic,
        font.underline,
    )


def text_signature(shape) -> tuple:
    if not getattr(shape, "has_text_frame", False):
        return ()
    return tuple(
        (
            paragraph.text,
            paragraph.level,
            str(paragraph.alignment),
            tuple(run_signature(run) for run in paragraph.runs if run.text),
        )
        for paragraph in shape.text_frame.paragraphs
    )


def footer_signature(presentation: Presentation, slide_number: int, expected_text: str) -> tuple | None:
    if slide_number < 1 or slide_number > len(presentation.slides):
        return None
    matches = [
        shape
        for shape in presentation.slides[slide_number - 1].shapes
        if getattr(shape, "has_text_frame", False) and normalize(shape.text) == expected_text
    ]
    if len(matches) != 1:
        return None
    shape = matches[0]
    return (
        shape.shape_id,
        shape.left,
        shape.top,
        shape.width,
        shape.height,
        str(shape.shape_type),
        text_signature(shape),
    )


def is_slide_number(shape) -> bool:
    if not getattr(shape, "is_placeholder", False):
        return False
    try:
        return "SLIDE_NUMBER" in str(shape.placeholder_format.type)
    except Exception:
        return False


def picture_signature(shape) -> tuple:
    if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
        return ()
    return (
        hashlib.sha256(shape.image.blob).hexdigest(),
        round(float(shape.crop_left), 6),
        round(float(shape.crop_right), 6),
        round(float(shape.crop_top), 6),
        round(float(shape.crop_bottom), 6),
    )


def table_signature(shape) -> tuple:
    if not shape.has_table:
        return ()
    table = shape.table
    return tuple(tuple(cell.text for cell in row.cells) for row in table.rows)


def slide_preserved(source_slide, candidate_slide) -> tuple[bool, list[str]]:
    source_shapes = {shape.shape_id: shape for shape in source_slide.shapes if not is_slide_number(shape)}
    candidate_shapes = {shape.shape_id: shape for shape in candidate_slide.shapes if not is_slide_number(shape)}
    errors = []
    if set(source_shapes) != set(candidate_shapes):
        errors.append(f"shape_ids {sorted(source_shapes)} != {sorted(candidate_shapes)}")
        return False, errors
    for shape_id, original in source_shapes.items():
        answer = candidate_shapes[shape_id]
        if str(original.shape_type) != str(answer.shape_type):
            errors.append(f"shape {shape_id} type")
            continue
        for name in ("left", "top", "width", "height"):
            if abs(int(getattr(original, name)) - int(getattr(answer, name))) > 15000:
                errors.append(f"shape {shape_id} {name}")
        if abs(float(original.rotation or 0) - float(answer.rotation or 0)) > 0.01:
            errors.append(f"shape {shape_id} rotation")
        if text_signature(original) != text_signature(answer):
            errors.append(f"shape {shape_id} text/style")
        if picture_signature(original) != picture_signature(answer):
            errors.append(f"shape {shape_id} picture")
        if table_signature(original) != table_signature(answer):
            errors.append(f"shape {shape_id} table")
    return not errors, errors


def original_deck_preserved(source: Presentation, candidate: Presentation) -> tuple[bool, dict]:
    if len(source.slides) != 23 or len(candidate.slides) != 24:
        return False, {"reason": "页数不匹配"}
    mappings = [(index, index) for index in range(22)] + [(22, 23)]
    errors = []
    for source_index, candidate_index in mappings:
        ok, slide_errors = slide_preserved(source.slides[source_index], candidate.slides[candidate_index])
        if not ok:
            errors.append({"source_slide": source_index + 1, "candidate_slide": candidate_index + 1, "errors": slide_errors[:8]})
    try:
        source_notes = [source.slides[index].notes_slide.notes_text_frame.text for index, _ in mappings]
        candidate_notes = [candidate.slides[index].notes_slide.notes_text_frame.text for _, index in mappings]
        notes_ok = source_notes == candidate_notes
    except Exception:
        notes_ok = False
    return not errors and notes_ok, {"slide_errors": errors[:8], "notes_ok": notes_ok}


def render(path: Path, outdir: Path) -> tuple[int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "profile"
    profile.mkdir()
    completed = subprocess.run(
        ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(path)],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=300,
        check=False,
    )
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return 0, completed.stdout[-500:]
    info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=30, check=False).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    return (int(match.group(1)) if match else 0), completed.stdout[-500:]


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {
        "criterion_id": identifier,
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
        "raw": {"description": description, "weight": weight, "type": kind, "passed": bool(passed)},
    }


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [item["criterion_id"] for item in criteria if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0 and item["raw"]["passed"])
    penalty = sum(abs(item["raw"]["weight"]) for item in criteria if item["raw"]["weight"] < 0 and not item["raw"]["passed"])
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [sha256(path)] if path.is_file() else [],
        "criterion_results": criteria,
        "reward": reward,
        "passed": reward >= 0.8 and not hurdles,
        "hurdle_failures": hurdles,
    }


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        return make_result(path, [result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配") for identifier in SPECS])
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        return make_result(path, [result_item(identifier, False, f"PPTX 打开失败：{exc}") for identifier in SPECS])

    criteria = []
    package_ok, package_evidence = package_checks(path)
    size_ok = len(candidate.slides) == 24 and (source.slide_width, source.slide_height) == (candidate.slide_width, candidate.slide_height)
    criteria.append(result_item("R001", package_ok and size_ok, f"页数={len(candidate.slides)}，尺寸一致={size_ok}，包={json.dumps(package_evidence, ensure_ascii=False)}"))

    bibliography_titles = []
    for index, slide in enumerate(candidate.slides):
        if any(normalize(shape.text) == "Bibliography" for shape in slide.shapes if getattr(shape, "has_text_frame", False)):
            bibliography_titles.append(index + 1)
    final_text = "\n".join(shape.text for shape in candidate.slides[-1].shapes if getattr(shape, "has_text_frame", False)) if candidate.slides else ""
    position_ok = bibliography_titles == [23] and "Thank you!" in final_text and "Questions?" in final_text
    criteria.append(result_item("R002", position_ok, f"Bibliography 页={bibliography_titles}，最后页文字={normalize(final_text)!r}"))

    records = bibliography_records(candidate)
    observed = tuple(value for value, _ in records)
    content_ok = Counter(observed) == Counter(EXPECTED_CITATIONS) and len(observed) == 9
    criteria.append(result_item("R003", content_ok, f"条数={len(observed)}，缺失={list((Counter(EXPECTED_CITATIONS)-Counter(observed)).elements())}，多余={list((Counter(observed)-Counter(EXPECTED_CITATIONS)).elements())}"))

    order_ok = observed == EXPECTED_CITATIONS and len({id(paragraph) for _, paragraph in records}) == 9
    criteria.append(result_item("R004", order_ok, f"前缀顺序={[value.split()[0] for value in observed]}，独立段落={len(records)}"))

    paragraph_map = {value: paragraph for value, paragraph in records}
    italic_detail = {citation.split()[0]: italic_substring(paragraph_map[citation], substring) if citation in paragraph_map else False for citation, substring in ITALIC_REQUIREMENTS.items()}
    italic_ok = all(italic_detail.values())
    criteria.append(result_item("R005", italic_ok, json.dumps(italic_detail, ensure_ascii=False)))

    footer_detail = []
    footer_ok = True
    for slide_number, expected in SOURCE_FOOTERS:
        source_signature = footer_signature(source, slide_number, expected)
        candidate_signature = footer_signature(candidate, slide_number, expected)
        equal = source_signature is not None and source_signature == candidate_signature
        footer_ok = footer_ok and equal
        footer_detail.append({"slide": slide_number, "prefix": expected[:18], "preserved": equal})
    criteria.append(result_item("R006", footer_ok, json.dumps(footer_detail, ensure_ascii=False)))

    originals_ok, originals_detail = original_deck_preserved(source, candidate)
    try:
        source_media = media_hashes(source_path)
        candidate_media = media_hashes(path)
        media_ok = source_media == candidate_media and sum(source_media.values()) == 15
    except Exception:
        media_ok = False
    criteria.append(result_item("R007", originals_ok and media_ok, f"原页={json.dumps(originals_detail, ensure_ascii=False)}，媒体={media_ok}"))

    bibliography_slide = candidate.slides[-2] if len(candidate.slides) >= 2 else None
    editable_shapes = [shape for shape in bibliography_slide.shapes if getattr(shape, "has_text_frame", False)] if bibliography_slide else []
    full_slide_pictures = [shape.shape_id for shape in bibliography_slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.width * shape.height >= candidate.slide_width * candidate.slide_height * 0.8] if bibliography_slide else []
    combined_text = "\n".join(shape.text for slide in candidate.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False)).lower()
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker in combined_text]
    editable_ok = len(editable_shapes) >= 2 and len(records) == 9 and not full_slide_pictures and not markers
    criteria.append(result_item("R008", editable_ok, f"文本对象={len(editable_shapes)}，整页图={full_slide_pictures}，标记={markers}"))

    pages, render_log = 0, ""
    try:
        with tempfile.TemporaryDirectory(prefix="pptarena083-render-") as temporary:
            pages, render_log = render(path, Path(temporary))
    except Exception as exc:
        render_log = f"渲染异常：{exc}"
    delivered = sorted(str(item.relative_to(artifact_dir)) for item in artifact_dir.rglob("*") if item.is_file())
    hygiene_ok = pages == 24 and delivered == [OUTPUT_NAME] and package_ok
    criteria.append(result_item("R009", hygiene_ok, f"渲染页数={pages}，文件={delivered}，日志={render_log}"))

    severe = not (position_ok and content_ok and order_ok and italic_ok and footer_ok and originals_ok and media_ok and editable_ok)
    criteria.append(result_item("R010", not severe, f"严重错误={severe}"))
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    path = args.artifact_dir / OUTPUT_NAME
    result = evaluate(path, args.artifact_dir) if path.is_file() else make_result(path, [result_item(identifier, False, f"缺失指定交付物：{OUTPUT_NAME}") for identifier in SPECS])
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{result['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
