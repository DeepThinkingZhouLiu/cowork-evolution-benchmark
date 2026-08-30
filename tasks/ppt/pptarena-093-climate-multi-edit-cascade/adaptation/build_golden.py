#!/usr/bin/env python3
"""PPTArena Case 93 上游答案的逐页审计与 Bibliography 纠错。"""

from __future__ import annotations

import argparse
import hashlib
import zipfile
from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt


INPUT_SHA256 = "14660dfd35f3ddb9aafe873159ca3871ecd2250d3e8f08db4e36cb61ca094fc7"
UPSTREAM_SHA256 = "5249026c8cc769dbf8586627c4d6d4acc8e7077843a4b2770dacd7617231bc4a"
SOURCE_URL = "https://climate.copernicus.eu/copernicus-2024-first-year-exceed-1.5degc-above-pre-industrial-level"
EXPECTED_MEDIA = {
    "1b9c7ff5b75f5db88aa7ca122a3a8e66e1315a9ef608114c1516c7a79a1ff9bc",
    "4b6f9528bf79830d7a43a9da7930766a7e90554f5ff887044533e386203c90ef",
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def media_hashes(path: Path) -> Counter:
    with zipfile.ZipFile(path) as archive:
        return Counter(
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("ppt/media/") and not name.endswith("/")
        )


def chart_contract(presentation: Presentation) -> bool:
    charts = [shape.chart for shape in presentation.slides[3].shapes if shape.shape_type == MSO_SHAPE_TYPE.CHART]
    if len(charts) != 1:
        return False
    chart = charts[0]
    categories = [category.label for category in chart.plots[0].categories]
    values = list(chart.series[0].values) if len(chart.series) == 1 else []
    return categories == ["1991–2020 Avg", "2023", "2024"] and all(abs(a - b) < 1e-6 for a, b in zip(values, [14.38, 14.98, 15.1], strict=False)) and len(values) == 3


def build(source: Path, upstream: Path, output: Path) -> None:
    if sha256(source) != INPUT_SHA256 or sha256(upstream) != UPSTREAM_SHA256:
        raise ValueError("冻结输入或上游 Ground Truth SHA256 不匹配")
    presentation = Presentation(upstream)
    if len(presentation.slides) != 7 or not chart_contract(presentation):
        raise ValueError("上游页数或第 4 页原生图表与已审计契约不匹配")
    if set(media_hashes(upstream)) != EXPECTED_MEDIA or sum(media_hashes(upstream).values()) != 2:
        raise ValueError("上游答案未恰好保留第 5 页的两张原图")

    final_slide = presentation.slides[-1]
    bodies = [
        shape
        for shape in final_slide.shapes
        if getattr(shape, "has_text_frame", False)
        and shape.text.strip()
        and shape.text.strip() != "Bibliography"
    ]
    if len(bodies) != 1:
        raise ValueError("上游 Bibliography 内容框不唯一")
    body = bodies[0]
    body.left = Inches(0.75)
    body.top = Inches(1.75)
    body.width = Inches(8.5)
    body.height = Inches(1.0)
    body.text_frame.clear()
    body.text_frame.word_wrap = True
    paragraph = body.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = SOURCE_URL
    run.font.size = Pt(16)
    run.font.underline = True
    run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
    run.hyperlink.address = SOURCE_URL

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    reopened = Presentation(output)
    final_runs = [run for shape in reopened.slides[-1].shapes if getattr(shape, "has_text_frame", False) for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.text == SOURCE_URL]
    if len(final_runs) != 1 or final_runs[0].hyperlink.address != SOURCE_URL or final_runs[0].font.underline is not True:
        raise ValueError("Bibliography 的单一来源超链接保存后不匹配")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--upstream", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.source, args.upstream, args.output)
    print(f"{args.output}\nsha256={sha256(args.output)}")


if __name__ == "__main__":
    main()
