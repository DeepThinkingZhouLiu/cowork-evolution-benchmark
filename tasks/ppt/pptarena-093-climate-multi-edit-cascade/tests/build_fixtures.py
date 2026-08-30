#!/usr/bin/env python3
"""PPTArena Case 93 专用 Judge 校准样例。"""

from __future__ import annotations

import argparse
import shutil
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Inches


OUTPUT_NAME = "climate_ground_truth.pptx"
SOURCE_URL = "https://climate.copernicus.eu/copernicus-2024-first-year-exceed-1.5degc-above-pre-industrial-level"
IMAGE_HASH_1 = "1b9c7ff5b75f5db88aa7ca122a3a8e66e1315a9ef608114c1516c7a79a1ff9bc"
CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"


def output_path(root: Path, case: str) -> Path:
    path = root / case / OUTPUT_NAME
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def save_mutation(golden: Path, output: Path, callback) -> None:
    presentation = Presentation(golden)
    callback(presentation)
    presentation.save(output)


def is_title(shape) -> bool:
    if not getattr(shape, "is_placeholder", False):
        return False
    try:
        return shape.placeholder_format.type in {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
    except Exception:
        return False


def fill_rgb(shape) -> str | None:
    try:
        return str(shape.fill.fore_color.rgb) if shape.fill.fore_color.rgb is not None else None
    except Exception:
        return None


def repackage(golden: Path, output: Path) -> None:
    with zipfile.ZipFile(golden) as reader, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_STORED) as writer:
        for info in reader.infolist():
            writer.writestr(info.filename, reader.read(info.filename))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--upstream", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    args = parser.parse_args()
    (args.outdir / "empty").mkdir(parents=True, exist_ok=True)
    output_path(args.outdir, "corrupt").write_bytes(b"not a valid pptx\n")
    shutil.copy2(args.source, output_path(args.outdir, "unmodified"))
    shutil.copy2(args.upstream, output_path(args.outdir, "upstream-unrefined"))
    repackage(args.golden, output_path(args.outdir, "valid-repackage"))

    def alt_theme(prs):
        for slide_number, slide in enumerate(prs.slides, 1):
            for shape in slide.shapes:
                if not getattr(shape, "has_text_frame", False):
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if not run.text.strip() or (slide_number == 7 and SOURCE_URL in run.text):
                            continue
                        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79) if is_title(shape) else RGBColor(0x44, 0x44, 0x44)
    save_mutation(args.golden, output_path(args.outdir, "valid-alt-theme"), alt_theme)

    def wrong_chart_value(prs):
        shape = next(shape for shape in prs.slides[3].shapes if shape.shape_type == MSO_SHAPE_TYPE.CHART)
        cache = shape.chart._chartSpace.find(f".//{{{CHART_NS}}}numCache")
        points = cache.findall(f"{{{CHART_NS}}}pt")
        points[0].find(f"{{{CHART_NS}}}v").text = "99"
    save_mutation(args.golden, output_path(args.outdir, "wrong-chart-value"), wrong_chart_value)

    def chart_removed(prs):
        shape = next(shape for shape in prs.slides[3].shapes if shape.shape_type == MSO_SHAPE_TYPE.CHART)
        shape._element.getparent().remove(shape._element)
    save_mutation(args.golden, output_path(args.outdir, "chart-removed"), chart_removed)

    def swapped_captions(prs):
        slide = prs.slides[4]
        captions = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False) and ("Global surface temperature" in shape.text or "Surface air temperature" in shape.text)]
        captions[0].text, captions[1].text = captions[1].text, captions[0].text
    save_mutation(args.golden, output_path(args.outdir, "swapped-captions"), swapped_captions)

    def wrong_progress(prs):
        candidates = [shape for shape in prs.slides[2].shapes if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and shape.top >= prs.slide_height * 0.9]
        foreground = next(shape for shape in candidates if fill_rgb(shape) == "E2671A")
        base = next(shape for shape in candidates if fill_rgb(shape) == "DDDDDD")
        foreground.width = base.width // 7
    save_mutation(args.golden, output_path(args.outdir, "wrong-progress-ratio"), wrong_progress)

    def missing_progress(prs):
        candidates = [shape for shape in prs.slides[1].shapes if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and shape.top >= prs.slide_height * 0.9]
        for shape in candidates:
            shape._element.getparent().remove(shape._element)
    save_mutation(args.golden, output_path(args.outdir, "missing-progress"), missing_progress)

    def extra_url(prs):
        body = next(shape for shape in prs.slides[-1].shapes if getattr(shape, "has_text_frame", False) and SOURCE_URL in shape.text)
        paragraph = body.text_frame.add_paragraph()
        paragraph.text = "https://example.com"
    save_mutation(args.golden, output_path(args.outdir, "extra-unverified-url"), extra_url)

    def wrong_link(prs):
        run = next(run for shape in prs.slides[-1].shapes if getattr(shape, "has_text_frame", False) for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if SOURCE_URL in run.text)
        run.hyperlink.address = "https://example.com"
    save_mutation(args.golden, output_path(args.outdir, "wrong-link-target"), wrong_link)

    def no_underline(prs):
        run = next(run for shape in prs.slides[-1].shapes if getattr(shape, "has_text_frame", False) for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if SOURCE_URL in run.text)
        run.font.underline = False
    save_mutation(args.golden, output_path(args.outdir, "link-not-underlined"), no_underline)

    def title_black(prs):
        title = next(shape for shape in prs.slides[1].shapes if is_title(shape))
        title.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 0, 0)
    save_mutation(args.golden, output_path(args.outdir, "wrong-title-color"), title_black)

    def body_red(prs):
        body = next(shape for shape in prs.slides[1].shapes if getattr(shape, "has_text_frame", False) and "In 2024" in shape.text)
        body.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x99, 0, 0)
    save_mutation(args.golden, output_path(args.outdir, "wrong-body-color"), body_red)

    def citation_removed(prs):
        shape = next(shape for shape in prs.slides[1].shapes if getattr(shape, "has_text_frame", False) and "[1]" in shape.text)
        shape.text = ""
    save_mutation(args.golden, output_path(args.outdir, "citation-removed"), citation_removed)

    def media_removed(prs):
        picture = next(shape for shape in prs.slides[4].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and __import__("hashlib").sha256(shape.image.blob).hexdigest() == IMAGE_HASH_1)
        picture._element.getparent().remove(picture._element)
    save_mutation(args.golden, output_path(args.outdir, "media-removed"), media_removed)

    placeholder = args.outdir / "flatten-placeholder.png"
    Image.new("RGB", (1200, 900), "white").save(placeholder)
    def flattened(prs):
        slide = prs.slides[4]
        for shape in list(slide.shapes):
            shape._element.getparent().remove(shape._element)
        slide.shapes.add_picture(str(placeholder), 0, 0, width=prs.slide_width, height=prs.slide_height)
    save_mutation(args.golden, output_path(args.outdir, "flattened-slide5"), flattened)

    def hardcoded(prs):
        box = prs.slides[-1].shapes.add_textbox(Inches(0.1), Inches(0.1), Inches(2.0), Inches(0.2))
        box.text = "solution/Golden"
    save_mutation(args.golden, output_path(args.outdir, "hardcoded-marker"), hardcoded)

    extra = args.outdir / "extra-output"
    extra.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.golden, extra / OUTPUT_NAME)
    (extra / "debug.txt").write_text("debug\n", encoding="utf-8")


if __name__ == "__main__":
    main()
