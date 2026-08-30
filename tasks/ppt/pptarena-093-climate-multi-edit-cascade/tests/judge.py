#!/usr/bin/env python3
"""PPTArena Case 93 的任务专属确定性 Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import posixpath
import re
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER


OUTPUT_NAME = "climate_ground_truth.pptx"
INPUT_NAME = "climate_original.pptx"
INPUT_SHA256 = "14660dfd35f3ddb9aafe873159ca3871ecd2250d3e8f08db4e36cb61ca094fc7"
TASK_ID = "cowork-evo/pptarena-093-climate-multi-edit-cascade"
RUBRIC_VERSION = "pptarena-case-093-adaptation-v0.1"
JUDGE_VERSION = "pptarena-case-093-deterministic-v0.1"
SOURCE_URL = "https://climate.copernicus.eu/copernicus-2024-first-year-exceed-1.5degc-above-pre-industrial-level"
STATIC_CHART_SHA = "230936ef4a19cd6e51a3de4accf4c572a9520821d4596fe3bb22f0bb971bbd16"
IMAGE_HASHES = {
    "Global surface temperature increase above pre-industrial (ERA5 & other sources).": "1b9c7ff5b75f5db88aa7ca122a3a8e66e1315a9ef608114c1516c7a79a1ff9bc",
    "Surface air temperature anomalies in 2024 relative to 1991–2020.": "4b6f9528bf79830d7a43a9da7930766a7e90554f5ff887044533e386203c90ef",
}
EXPECTED_CATEGORIES = ["1991–2020 Avg", "2023", "2024"]
EXPECTED_VALUES = [14.38, 14.98, 15.10]
SUSPICIOUS_MARKERS = ["solution/golden", "judge-result", "artifact_dir", "pptarena-case-093", "/root/"]

SPECS = {
    "R001": ("PPTX 完整、保留画布且最终 7 页", 6, "hurdle"),
    "R002": ("7 页深蓝标题与深灰正文主题一致", 12, "hurdle"),
    "R003": ("第 4 页原生簇状柱形图的类别、数值和排版正确", 20, "hurdle"),
    "R004": ("第 5 页两张原图与各自图注关联清晰、无重叠", 14, "hurdle"),
    "R005": ("7 页底部进度条按 1/7 到 7/7 递增", 14, "hurdle"),
    "R006": ("末页只有单一准确 Copernicus URL 可点击超链接", 14, "hurdle"),
    "R007": ("原 [1]–[6]、链接目标和全部文字事实保留", 8, "hurdle"),
    "R008": ("静态图表删除、仅保留两张原图且内容原生可编辑", 8, "hurdle"),
    "R009": ("7 页可渲染且输出目录干净", 4, "positive"),
    "R010": ("无主题、图表、图注、进度、URL、引用或压平严重错误", -20, "penalty"),
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("\u00a0", " ").replace("\v", " ")).strip()


def locate_source() -> Path | None:
    candidates = [Path("/app/input_files") / INPUT_NAME, Path(__file__).resolve().parents[1] / "data/input_files" / INPUT_NAME]
    return next((path for path in candidates if path.is_file()), None)


def package_checks(path: Path) -> tuple[bool, dict]:
    evidence = {"broken_member": None, "missing_targets": [], "macros": []}
    try:
        with zipfile.ZipFile(path) as archive:
            evidence["broken_member"] = archive.testzip()
            names = set(archive.namelist())
            evidence["macros"] = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
            for rels_name in sorted(name for name in names if name.endswith(".rels")):
                root = etree.fromstring(archive.read(rels_name))
                if rels_name == "_rels/.rels":
                    owner = ""
                else:
                    prefix, filename = rels_name.split("/_rels/", 1)
                    owner = posixpath.join(prefix, filename[:-5])
                base = posixpath.dirname(owner)
                for relation in root:
                    if relation.attrib.get("TargetMode") == "External":
                        continue
                    target = relation.attrib.get("Target", "")
                    resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                    if resolved not in names:
                        evidence["missing_targets"].append(f"{rels_name} -> {resolved}")
        return evidence["broken_member"] is None and not evidence["missing_targets"] and not evidence["macros"], evidence
    except Exception as exc:
        evidence["exception"] = str(exc)
        return False, evidence


def media_hashes(path: Path) -> Counter:
    with zipfile.ZipFile(path) as archive:
        return Counter(
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("ppt/media/") and not name.endswith("/")
        )


def rgb_tuple(run) -> tuple[int, int, int] | None:
    try:
        rgb = run.font.color.rgb
        if rgb is None:
            return None
        value = str(rgb)
        return tuple(int(value[index:index + 2], 16) for index in (0, 2, 4))
    except Exception:
        return None


def is_dark_blue(rgb: tuple[int, int, int] | None) -> bool:
    if rgb is None:
        return False
    red, green, blue = rgb
    return max(rgb) <= 150 and blue >= green + 20 and blue >= red + 30 and blue >= 60


def is_dark_gray(rgb: tuple[int, int, int] | None) -> bool:
    if rgb is None:
        return False
    return max(rgb) <= 100 and max(rgb) - min(rgb) <= 18


def is_title_shape(shape) -> bool:
    if not getattr(shape, "is_placeholder", False):
        return False
    try:
        return shape.placeholder_format.type in {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
    except Exception:
        return False


def theme_analysis(presentation: Presentation) -> tuple[bool, dict]:
    title_colors = []
    body_colors = []
    failures = []
    for slide_number, slide in enumerate(presentation.slides, 1):
        titles = [shape for shape in slide.shapes if is_title_shape(shape) and normalize(shape.text)]
        if len(titles) != 1:
            failures.append(f"slide {slide_number} title_count={len(titles)}")
        for shape in titles:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip():
                        color = rgb_tuple(run)
                        title_colors.append(color)
                        if not is_dark_blue(color):
                            failures.append(f"slide {slide_number} title_color={color}")
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False) or is_title_shape(shape):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if not run.text.strip():
                        continue
                    if slide_number == 7 and SOURCE_URL in run.text:
                        continue
                    color = rgb_tuple(run)
                    body_colors.append(color)
                    if not is_dark_gray(color):
                        failures.append(f"slide {slide_number} body_color={color} text={run.text[:20]!r}")
    consistent = len(set(title_colors)) == 1 and len(set(body_colors)) == 1
    return not failures and consistent and len(title_colors) >= 7 and bool(body_colors), {
        "title_colors": sorted(map(str, set(title_colors))),
        "body_colors": sorted(map(str, set(body_colors))),
        "consistent": consistent,
        "failures": failures[:10],
    }


def overlap_area(first, second) -> int:
    left = max(first.left, second.left)
    top = max(first.top, second.top)
    right = min(first.left + first.width, second.left + second.width)
    bottom = min(first.top + first.height, second.top + second.height)
    return max(0, right - left) * max(0, bottom - top)


def within_canvas(shape, presentation: Presentation) -> bool:
    return shape.left >= 0 and shape.top >= 0 and shape.left + shape.width <= presentation.slide_width and shape.top + shape.height <= presentation.slide_height


def chart_analysis(presentation: Presentation) -> tuple[bool, dict]:
    if len(presentation.slides) < 4:
        return False, {"reason": "缺少第 4 页"}
    slide = presentation.slides[3]
    chart_shapes = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.CHART]
    pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    descriptions = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False) and "This bar chart compares" in shape.text]
    if len(chart_shapes) != 1 or len(descriptions) != 1:
        return False, {"charts": len(chart_shapes), "descriptions": len(descriptions), "pictures": len(pictures)}
    shape = chart_shapes[0]
    chart = shape.chart
    categories = [category.label for category in chart.plots[0].categories]
    series = list(chart.series)
    values = list(series[0].values) if len(series) == 1 else []
    numeric_ok = categories == EXPECTED_CATEGORIES and len(values) == 3 and all(math.isclose(float(value), expected, abs_tol=1e-6) for value, expected in zip(values, EXPECTED_VALUES, strict=False))
    column_ok = "COLUMN" in str(chart.chart_type)
    layout_ok = overlap_area(shape, descriptions[0]) == 0 and within_canvas(shape, presentation) and within_canvas(descriptions[0], presentation)
    return numeric_ok and column_ok and not pictures and layout_ok, {
        "categories": categories,
        "values": values,
        "type": str(chart.chart_type),
        "pictures": len(pictures),
        "overlap": overlap_area(shape, descriptions[0]),
        "layout_ok": layout_ok,
    }


def picture_hash(shape) -> str:
    return hashlib.sha256(shape.image.blob).hexdigest()


def adjacent(first, second) -> bool:
    if overlap_area(first, second):
        return False
    x_overlap = max(0, min(first.left + first.width, second.left + second.width) - max(first.left, second.left))
    y_overlap = max(0, min(first.top + first.height, second.top + second.height) - max(first.top, second.top))
    x_gap = max(0, max(first.left, second.left) - min(first.left + first.width, second.left + second.width))
    y_gap = max(0, max(first.top, second.top) - min(first.top + first.height, second.top + second.height))
    horizontal = y_overlap >= min(first.height, second.height) * 0.25 and x_gap <= 457200
    vertical = x_overlap >= min(first.width, second.width) * 0.25 and y_gap <= 457200
    return horizontal or vertical


def slide5_analysis(presentation: Presentation) -> tuple[bool, dict]:
    if len(presentation.slides) < 5:
        return False, {"reason": "缺少第 5 页"}
    slide = presentation.slides[4]
    pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    by_hash = {picture_hash(shape): shape for shape in pictures}
    detail = {"picture_hashes": sorted(by_hash), "associations": {}, "overlaps": []}
    ok = len(pictures) == 2 and set(by_hash) == set(IMAGE_HASHES.values())
    all_elements = list(pictures)
    for caption_text, expected_hash in IMAGE_HASHES.items():
        captions = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False) and normalize(shape.text) == caption_text]
        caption = captions[0] if len(captions) == 1 else None
        own = by_hash.get(expected_hash)
        other = next((shape for value, shape in by_hash.items() if value != expected_hash), None)
        association = caption is not None and own is not None and adjacent(caption, own) and (other is None or not adjacent(caption, other))
        detail["associations"][caption_text[:28]] = association
        ok = ok and association
        if caption is not None:
            all_elements.append(caption)
    for index, first in enumerate(all_elements):
        if not within_canvas(first, presentation):
            ok = False
        for second in all_elements[index + 1:]:
            if overlap_area(first, second):
                detail["overlaps"].append([first.shape_id, second.shape_id])
                ok = False
    return ok, detail


def fill_rgb(shape) -> str | None:
    try:
        return str(shape.fill.fore_color.rgb) if shape.fill.fore_color.rgb is not None else None
    except Exception:
        return None


def progress_analysis(presentation: Presentation) -> tuple[bool, dict]:
    if len(presentation.slides) != 7:
        return False, {"reason": "页数不是 7"}
    per_slide = []
    base_color = None
    fill_color = None
    base_geometry = None
    ok = True
    for index, slide in enumerate(presentation.slides, 1):
        candidates = [
            shape for shape in slide.shapes
            if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
            and shape.top >= presentation.slide_height * 0.90
            and 45000 <= shape.height <= 300000
            and shape.width >= presentation.slide_width * 0.10
        ]
        if index == 1 and len(candidates) == 2:
            base = max(candidates, key=lambda shape: shape.width)
            fill = min(candidates, key=lambda shape: shape.width)
            base_color, fill_color = fill_rgb(base), fill_rgb(fill)
            base_geometry = (base.left, base.top, base.width, base.height)
        else:
            base_matches = [shape for shape in candidates if fill_rgb(shape) == base_color]
            fill_matches = [shape for shape in candidates if fill_rgb(shape) == fill_color]
            base = base_matches[0] if len(base_matches) == 1 else None
            fill = fill_matches[0] if len(fill_matches) == 1 else None
        if base is None or fill is None or base_geometry is None:
            per_slide.append({"slide": index, "candidates": len(candidates), "ok": False})
            ok = False
            continue
        expected = index / 7
        ratio = fill.width / base.width if base.width else 0
        geometry_ok = all(abs(value - target) <= 45720 for value, target in zip((base.left, base.top, base.width, base.height), base_geometry, strict=True))
        slide_ok = abs(ratio - expected) <= 0.025 and geometry_ok and abs(fill.left - base.left) <= 9144 and abs(fill.top - base.top) <= 9144 and fill.height == base.height and within_canvas(base, presentation) and within_canvas(fill, presentation)
        per_slide.append({"slide": index, "ratio": round(ratio, 4), "expected": round(expected, 4), "geometry_ok": geometry_ok, "ok": slide_ok})
        ok = ok and slide_ok
    return ok, {"base_color": base_color, "fill_color": fill_color, "slides": per_slide}


def bibliography_analysis(presentation: Presentation) -> tuple[bool, dict]:
    if len(presentation.slides) != 7:
        return False, {"reason": "页数不是 7"}
    slide = presentation.slides[-1]
    titles = [shape for shape in slide.shapes if is_title_shape(shape) and normalize(shape.text) == "Bibliography"]
    runs = [run for shape in slide.shapes if getattr(shape, "has_text_frame", False) for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.text.strip()]
    url_runs = [run for run in runs if normalize(run.text) == SOURCE_URL]
    visible_texts = [normalize(run.text) for run in runs if normalize(run.text) != "Bibliography"]
    hyperlink_targets = [run.hyperlink.address for run in url_runs]
    link_ok = len(url_runs) == 1 and hyperlink_targets == [SOURCE_URL] and url_runs[0].font.underline is True and rgb_tuple(url_runs[0]) is not None and not is_dark_gray(rgb_tuple(url_runs[0]))
    pictures = [shape.shape_id for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    ok = len(titles) == 1 and visible_texts == [SOURCE_URL] and link_ok and not pictures
    return ok, {"titles": len(titles), "visible_texts": visible_texts, "targets": hyperlink_targets, "underlined": [run.font.underline for run in url_runs], "pictures": pictures}


def text_counter(presentation: Presentation, slide_count: int) -> Counter:
    return Counter(
        normalize(shape.text)
        for slide in list(presentation.slides)[:slide_count]
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and normalize(shape.text)
    )


def citation_signature(presentation: Presentation) -> dict:
    result = {}
    for slide_number in range(2, 7):
        for shape in presentation.slides[slide_number - 1].shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = normalize(shape.text)
            if re.fullmatch(r"\[(?:\d+)\](?: \[(?:\d+)\])?", text):
                targets = tuple(run.hyperlink.address for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.text.strip())
                result[slide_number] = (text, targets)
    return result


def render(path: Path, outdir: Path) -> tuple[int, str]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "profile"
    profile.mkdir()
    completed = subprocess.run(["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(path)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300, check=False)
    pdf = outdir / f"{path.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        return 0, completed.stdout[-500:]
    info = subprocess.run(["pdfinfo", str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=30, check=False).stdout
    match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
    return (int(match.group(1)) if match else 0), completed.stdout[-500:]


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {"criterion_id": identifier, "score": 1.0 if passed else 0.0, "evidence": evidence, "method": "deterministic", "raw": {"description": description, "weight": weight, "type": kind, "passed": bool(passed)}}


def make_result(path: Path, criteria: list[dict]) -> dict:
    hurdles = [item["criterion_id"] for item in criteria if item["raw"]["type"] == "hurdle" and not item["raw"]["passed"]]
    total = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0)
    earned = sum(item["raw"]["weight"] for item in criteria if item["raw"]["weight"] > 0 and item["raw"]["passed"])
    penalty = sum(abs(item["raw"]["weight"]) for item in criteria if item["raw"]["weight"] < 0 and not item["raw"]["passed"])
    reward = 0.0 if hurdles else max(0.0, min(1.0, (earned - penalty) / total))
    return {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [sha256(path)] if path.is_file() else [], "criterion_results": criteria, "reward": reward, "passed": reward >= 0.8 and not hurdles, "hurdle_failures": hurdles}


def evaluate(path: Path, artifact_dir: Path) -> dict:
    source_path = locate_source()
    if source_path is None or sha256(source_path) != INPUT_SHA256:
        return make_result(path, [result_item(identifier, False, "冻结输入缺失或 SHA256 不匹配") for identifier in SPECS])
    try:
        source = Presentation(source_path)
        candidate = Presentation(path)
    except Exception as exc:
        return make_result(path, [result_item(identifier, False, f"PPTX 打开失败：{exc}") for identifier in SPECS])

    criteria = []
    package_ok, package_evidence = package_checks(path)
    size_ok = len(candidate.slides) == 7 and (source.slide_width, source.slide_height) == (candidate.slide_width, candidate.slide_height)
    criteria.append(result_item("R001", package_ok and size_ok, f"页数={len(candidate.slides)}，尺寸一致={size_ok}，包={json.dumps(package_evidence, ensure_ascii=False)}"))

    theme_ok, theme_detail = theme_analysis(candidate)
    criteria.append(result_item("R002", theme_ok, json.dumps(theme_detail, ensure_ascii=False)))

    chart_ok, chart_detail = chart_analysis(candidate)
    criteria.append(result_item("R003", chart_ok, json.dumps(chart_detail, ensure_ascii=False)))

    slide5_ok, slide5_detail = slide5_analysis(candidate)
    criteria.append(result_item("R004", slide5_ok, json.dumps(slide5_detail, ensure_ascii=False)))

    progress_ok, progress_detail = progress_analysis(candidate)
    criteria.append(result_item("R005", progress_ok, json.dumps(progress_detail, ensure_ascii=False)))

    bibliography_ok, bibliography_detail = bibliography_analysis(candidate)
    criteria.append(result_item("R006", bibliography_ok, json.dumps(bibliography_detail, ensure_ascii=False)))

    source_texts = text_counter(source, 6)
    candidate_texts = text_counter(candidate, 6)
    source_citations = citation_signature(source)
    candidate_citations = citation_signature(candidate)
    preservation_ok = source_texts == candidate_texts and source_citations == candidate_citations and set(source_citations) == {2, 3, 4, 5, 6}
    criteria.append(result_item("R007", preservation_ok, f"文字一致={source_texts == candidate_texts}，引用源={json.dumps({str(k): v for k, v in source_citations.items()}, ensure_ascii=False)}，引用候选={json.dumps({str(k): v for k, v in candidate_citations.items()}, ensure_ascii=False)}"))

    try:
        media = media_hashes(path)
        media_ok = media == Counter(IMAGE_HASHES.values()) and STATIC_CHART_SHA not in media
    except Exception:
        media, media_ok = Counter(), False
    full_slide_pictures = [shape.shape_id for slide in candidate.slides for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.width * shape.height >= candidate.slide_width * candidate.slide_height * 0.8]
    combined_text = "\n".join(shape.text for slide in candidate.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False)).lower()
    markers = [marker for marker in SUSPICIOUS_MARKERS if marker in combined_text]
    editable_ok = media_ok and chart_ok and not full_slide_pictures and not markers
    criteria.append(result_item("R008", editable_ok, f"媒体={dict(media)}，静态图表存在={STATIC_CHART_SHA in media}，整页图={full_slide_pictures}，标记={markers}"))

    pages, render_log = 0, ""
    try:
        with tempfile.TemporaryDirectory(prefix="pptarena093-render-") as temporary:
            pages, render_log = render(path, Path(temporary))
    except Exception as exc:
        render_log = f"渲染异常：{exc}"
    delivered = sorted(str(item.relative_to(artifact_dir)) for item in artifact_dir.rglob("*") if item.is_file())
    hygiene_ok = pages == 7 and delivered == [OUTPUT_NAME] and package_ok
    criteria.append(result_item("R009", hygiene_ok, f"渲染页数={pages}，文件={delivered}，日志={render_log}"))

    severe = not (theme_ok and chart_ok and slide5_ok and progress_ok and bibliography_ok and preservation_ok and editable_ok)
    criteria.append(result_item("R010", not severe, f"严重错误={severe}"))
    return make_result(path, criteria)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    path = args.artifact_dir / OUTPUT_NAME
    result = evaluate(path, args.artifact_dir) if path.is_file() else make_result(path, [result_item(identifier, False, f"缺失指定交付物：{OUTPUT_NAME}") for identifier in SPECS])
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{result['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
