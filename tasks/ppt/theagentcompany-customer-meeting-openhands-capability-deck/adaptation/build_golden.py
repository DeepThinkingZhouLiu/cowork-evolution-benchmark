#!/usr/bin/env python3
"""TheAgentCompany 客户会前 OpenHands 能力演示题的逐题 Golden。"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HASHES={"Original_Request.md":"eb8a006788c291de4c76a47e8f72a33670395d1f2999c99b45c7e0dd8c5bbdef","RocketChat_Recipient_Context.json":"865c354cfc784ef51994a0b0159dd396f3a76c3f8ba01046a039d3d0d1ac07b8","OpenHands_Evaluation_Snapshot.json":"aa411f23fbb791720b0c9274a5900685c13d929f2dce1c20d711a406a26d3338"}
NAVY=RGBColor(0x0C,0x1B,0x2A); BLUE=RGBColor(0x31,0x74,0xD8); CYAN=RGBColor(0x16,0xAC,0xBE); VIOLET=RGBColor(0x7A,0x5A,0xD6); GREEN=RGBColor(0x17,0x91,0x69); AMBER=RGBColor(0xDF,0x8B,0x15); WHITE=RGBColor(0xFF,0xFF,0xFF); INK=RGBColor(0x16,0x22,0x35); MUTED=RGBColor(0x66,0x73,0x84); LINE=RGBColor(0xDF,0xE6,0xED); SURFACE=RGBColor(0xF4,0xF7,0xFA); PALE=RGBColor(0xE8,0xF1,0xFB)

def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def text(slide,value,left,top,width,height,*,size=14,color=INK,bold=False,align=PP_ALIGN.LEFT,margin=.04,valign=MSO_ANCHOR.TOP):
    shape=slide.shapes.add_textbox(Inches(left),Inches(top),Inches(width),Inches(height)); frame=shape.text_frame; frame.clear(); frame.word_wrap=True; frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=Inches(margin); frame.vertical_anchor=valign; paragraph=frame.paragraphs[0]; paragraph.alignment=align; run=paragraph.add_run(); run.text=str(value); run.font.name="Aptos"; run.font.size=Pt(size); run.font.color.rgb=color; run.font.bold=bold; return shape
def rect(slide,left,top,width,height,fill,line=None,rounded=True):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(left),Inches(top),Inches(width),Inches(height)); shape.fill.solid(); shape.fill.fore_color.rgb=fill; shape.line.color.rgb=line or fill; return shape
def bg(slide,color): slide.background.fill.solid(); slide.background.fill.fore_color.rgb=color
def footer(slide,page,dark=False):
    color=RGBColor(0xA9,0xB7,0xC8) if dark else MUTED; text(slide,"Source: OpenHands_Evaluation_Snapshot.json  ·  upstream evaluator @ 98b68ef8",.65,7.12,11.2,.18,size=7.2,color=color); text(slide,str(page),12.2,7.12,.35,.18,size=8,color=color,align=PP_ALIGN.RIGHT)
def verify(root):
    for name,digest in HASHES.items():
        path=root/name
        if not path.is_file() or sha(path)!=digest: raise ValueError(f"冻结输入错误：{name}")
    snapshot=json.loads((root/"OpenHands_Evaluation_Snapshot.json").read_text()); categories=snapshot["categories"]
    if [item["order"] for item in categories]!=[1,2,3] or [len(item["benchmarks"]) for item in categories]!=[8,2,6] or sum(len(item["benchmarks"]) for item in categories)!=16: raise ValueError("评测快照不一致")
    recipient=json.loads((root/"RocketChat_Recipient_Context.json").read_text())
    if "Huang Jie" not in recipient: raise ValueError("收件人上下文不一致")
    return snapshot
def add_chip(slide,value,left,top,width,accent,index):
    rect(slide,left,top,width,.78,WHITE,LINE); rect(slide,left,top,.06,.78,accent,accent,rounded=False); text(slide,f"{index:02d}",left+.2,top+.16,.38,.22,size=8,color=accent,bold=True); text(slide,value,left+.62,top+.14,width-.78,.31,size=15,color=NAVY,bold=True); text(slide,"EVALUATION ENTRY",left+.62,top+.49,width-.78,.14,size=6.7,color=MUTED,bold=True)
def domain_slide(prs,blank,category,page,accent):
    slide=prs.slides.add_slide(blank); bg(slide,SURFACE)
    text(slide,f"DOMAIN {category['order']:02d}  /  FROZEN INVENTORY",.68,.34,5.8,.2,size=8.7,color=accent,bold=True); text(slide,category["name"],.68,.72,9.5,.56,size=30,color=NAVY,bold=True); text(slide,f"UPSTREAM EVALUATOR ANCHOR  ·  {category['upstream_evaluator_anchor']}",.7,1.38,6.2,.2,size=8.2,color=MUTED,bold=True); rect(slide,.68,1.78,12.0,.02,LINE,rounded=False)
    count=len(category["benchmarks"]); cols=2 if count>2 else 1; chip_width=5.72 if cols==2 else 8.3; start_left=.72 if cols==2 else 2.48; rows=(count+cols-1)//cols
    for index,name in enumerate(category["benchmarks"],1):
        col=(index-1)%2 if cols==2 else 0; row=(index-1)//cols; left=start_left+col*6.02 if cols==2 else start_left; top=2.18+row*.93; add_chip(slide,name,left,top,chip_width,accent,index)
    bottom=2.18+rows*.93+.2; rect(slide,.72,bottom,11.86,.74,PALE,PALE); text(slide,f"What this establishes: {count} benchmark entries are grouped under {category['name']} by the upstream evaluator.",.98,bottom+.16,11.25,.2,size=10.5,color=BLUE,bold=True); text(slide,"What it does not establish: pass rate, model version, score, ranking, or customer outcome.",.98,bottom+.43,11.25,.18,size=8.2,color=MUTED); footer(slide,page)
def build(input_dir,output):
    snapshot=verify(input_dir); prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]
    slide=prs.slides.add_slide(blank); bg(slide,NAVY); text(slide,"CUSTOMER MEETING BRIEF  /  OPENHANDS",.72,.65,6.2,.22,size=9.5,color=CYAN,bold=True); text(slide,"OpenHands Agent\nCapability Inventory",.72,1.2,7.7,1.24,size=35,color=WHITE,bold=True); text(slide,"A source-bounded view of the benchmark evaluation surface.",.75,2.79,7.2,.3,size=14,color=RGBColor(0xC5,0xD1,0xDF)); text(slide,"Potential customer discussion  ·  frozen GitLab evaluation tree",.75,3.32,7.7,.24,size=10,color=CYAN,bold=True)
    for index,(value,label,accent) in enumerate([("3","task domains",BLUE),("16","benchmark entries",GREEN),("0","performance claims",AMBER)]):
        left=.72+index*4.08; rect(slide,left,4.48,3.72,1.35,RGBColor(0x17,0x2B,0x43),RGBColor(0x2F,0x47,0x67)); rect(slide,left,4.48,.06,1.35,accent,accent,rounded=False); text(slide,value,left+.25,4.69,3.0,.38,size=24,color=WHITE,bold=True); text(slide,label.upper(),left+.25,5.3,3.0,.2,size=8.3,color=accent,bold=True)
    text(slide,"Inventory coverage is not measured performance. No score, success rate, model version, or comparative ranking is present in the source.",.75,6.33,11.5,.34,size=9,color=RGBColor(0xB9,0xC5,0xD5)); footer(slide,1,True)

    slide=prs.slides.add_slide(blank); bg(slide,SURFACE); text(slide,"CAPABILITY MAP",.68,.34,3.1,.2,size=8.7,color=BLUE,bold=True); text(slide,"Three domains. One evidence boundary.",.68,.72,10.8,.54,size=29,color=NAVY,bold=True); text(slide,"The source supplies category membership and names—not performance evidence.",.7,1.42,8.6,.24,size=11,color=MUTED); rect(slide,.68,1.82,12,.02,LINE,rounded=False)
    colors=[BLUE,CYAN,VIOLET]
    for index,category in enumerate(snapshot["categories"]):
        left=.72+index*4.03; rect(slide,left,2.25,3.76,2.78,WHITE,LINE); rect(slide,left,2.25,3.76,.08,colors[index],colors[index],rounded=False); text(slide,f"0{category['order']}",left+.26,2.57,.5,.24,size=10,color=colors[index],bold=True); text(slide,category["name"],left+.26,3.05,3.2,.58,size=20,color=NAVY,bold=True); text(slide,str(len(category["benchmarks"])),left+.26,3.92,1.0,.42,size=27,color=colors[index],bold=True); text(slide,"BENCHMARK ENTRIES",left+1.13,4.06,2.25,.2,size=7.5,color=MUTED,bold=True)
    rect(slide,.72,5.47,11.86,.94,PALE,PALE); text(slide,"Customer-safe statement",.98,5.67,2.1,.2,size=9,color=BLUE,bold=True); text(slide,"OpenHands maintains evaluation entries across Software Engineering, Web Browsing, and Misc. Assistance in the frozen repository snapshot.",3.0,5.63,9.1,.4,size=11,color=NAVY,bold=True); footer(slide,2)

    domain_slide(prs,blank,snapshot["categories"][0],3,BLUE); domain_slide(prs,blank,snapshot["categories"][1],4,CYAN); domain_slide(prs,blank,snapshot["categories"][2],5,VIOLET)

    slide=prs.slides.add_slide(blank); bg(slide,NAVY); text(slide,"EVIDENCE & HANDOFF",.72,.62,4.2,.2,size=9,color=CYAN,bold=True); text(slide,"Keep the claim\ninside the source.",.72,1.08,6.3,1.08,size=31,color=WHITE,bold=True); text(slide,"Evaluation source",.75,2.65,2.0,.2,size=8.5,color=CYAN,bold=True); text(slide,snapshot["evaluation_page_url"],.75,3.05,11.5,.3,size=12,color=WHITE,bold=True)
    rect(slide,.72,3.7,7.6,1.63,RGBColor(0x17,0x2B,0x43),RGBColor(0x2F,0x47,0x67)); text(slide,"Evidence boundary",.98,3.98,2.1,.2,size=9,color=AMBER,bold=True); text(slide,"Names + domain membership only. No scores, success rates, versions, rankings, or customer outcomes.",.98,4.39,6.95,.56,size=13,color=WHITE,bold=True)
    rect(slide,8.62,3.7,3.96,1.63,RGBColor(0x17,0x2B,0x43),RGBColor(0x2F,0x47,0x67)); text(slide,"Offline handoff",8.9,3.98,2.2,.2,size=9,color=GREEN,bold=True); text(slide,"openhands_intro.pptx",8.9,4.36,3.2,.25,size=12,color=WHITE,bold=True); text(slide,"Huang Jie  ·  RocketChat DM",8.9,4.73,3.2,.2,size=8.5,color=RGBColor(0xBD,0xC9,0xD8)); text(slide,"No live network side effect is claimed.",.75,5.83,8.0,.25,size=10,color=CYAN,bold=True); text(slide,"Delivery status: not_executed_offline",.75,6.22,6.0,.23,size=9,color=RGBColor(0xB9,0xC5,0xD5)); footer(slide,6,True)
    output.parent.mkdir(parents=True,exist_ok=True); prs.save(output)
    if len(Presentation(output).slides)!=6: raise ValueError("页数不匹配")
    print(f"{output}\nsha256={sha(output)}")

def main():
    parser=argparse.ArgumentParser(); parser.add_argument("--input-dir",type=Path,required=True); parser.add_argument("--output",type=Path,required=True); args=parser.parse_args(); build(args.input_dir,args.output)

if __name__=="__main__": main()
