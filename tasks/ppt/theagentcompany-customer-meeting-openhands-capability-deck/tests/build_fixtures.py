#!/usr/bin/env python3
from __future__ import annotations
import argparse,json,shutil,zipfile
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

PPTX_NAME="openhands_intro.pptx"; JSON_NAME="Delivery_Record.json"
def directory(root,name): path=root/name; path.mkdir(parents=True,exist_ok=True); return path
def replace_on_slide(slide,old,new):
    count=0
    for shape in slide.shapes:
        if not getattr(shape,"has_text_frame",False): continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if old in run.text: run.text=run.text.replace(old,new); count+=1
    if not count: raise ValueError(old)
def replace_all(prs,old,new):
    count=0
    for slide in prs.slides:
        for shape in slide.shapes:
            if not getattr(shape,"has_text_frame",False): continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if old in run.text: run.text=run.text.replace(old,new); count+=1
    if not count: raise ValueError(old)
def mutate(golden,record,target,operation):
    target=directory(target.parent,target.name); prs=Presentation(golden); operation(prs); prs.save(target/PPTX_NAME); shutil.copy2(record,target/JSON_NAME)
def mutate_json(golden,record,target,key,value):
    target=directory(target.parent,target.name); shutil.copy2(golden,target/PPTX_NAME); data=json.loads(record.read_text()); data[key]=value; (target/JSON_NAME).write_text(json.dumps(data,ensure_ascii=False,indent=2)+"\n")
def main():
    parser=argparse.ArgumentParser(); parser.add_argument("--golden-dir",type=Path,required=True); parser.add_argument("--outdir",type=Path,required=True); args=parser.parse_args(); golden=args.golden_dir/PPTX_NAME; record=args.golden_dir/JSON_NAME
    directory(args.outdir,"empty"); output=directory(args.outdir,"corrupt"); (output/PPTX_NAME).write_bytes(b"bad\n"); shutil.copy2(record,output/JSON_NAME)
    output=directory(args.outdir,"ppt-only"); shutil.copy2(golden,output/PPTX_NAME); output=directory(args.outdir,"json-only"); shutil.copy2(record,output/JSON_NAME)
    output=directory(args.outdir,"valid-repackage")
    with zipfile.ZipFile(golden) as source,zipfile.ZipFile(output/PPTX_NAME,"w",compression=zipfile.ZIP_STORED) as target:
        for info in source.infolist(): target.writestr(info.filename,source.read(info.filename))
    shutil.copy2(record,output/JSON_NAME)
    def alt_cover(prs): prs.slides[0].background.fill.solid(); prs.slides[0].background.fill.fore_color.rgb=RGBColor(0x12,0x2A,0x35)
    mutate(golden,record,args.outdir/"valid-alt-cover",alt_cover)
    mutate(golden,record,args.outdir/"missing-software-benchmark",lambda prs:replace_on_slide(prs.slides[2],"AiderBench","Entry omitted"))
    mutate(golden,record,args.outdir/"missing-web-benchmark",lambda prs:replace_on_slide(prs.slides[3],"MiniWoB","Entry omitted"))
    mutate(golden,record,args.outdir/"missing-misc-benchmark",lambda prs:replace_on_slide(prs.slides[4],"ProofWriter","Entry omitted"))
    mutate(golden,record,args.outdir/"foreign-domain-entry",lambda prs:replace_on_slide(prs.slides[3],"MiniWoB","SWE-Bench"))
    def unsupported_claim(prs):
        box=prs.slides[0].shapes.add_textbox(Inches(.8),Inches(4.15),Inches(6),Inches(.25)); box.text="OpenHands achieves a 95% pass rate."
    mutate(golden,record,args.outdir/"unsupported-performance-claim",unsupported_claim)
    mutate(golden,record,args.outdir/"missing-evidence-boundary",lambda prs:replace_all(prs,"What it does not establish:","Customer proof:"))
    mutate_json(golden,record,args.outdir/"wrong-recipient","recipient","Li Lei")
    mutate_json(golden,record,args.outdir/"wrong-channel","channel","Email")
    mutate_json(golden,record,args.outdir/"false-network-success","network_side_effect","executed_successfully")
    mutate(golden,record,args.outdir/"cross-artifact-recipient",lambda prs:replace_on_slide(prs.slides[5],"Huang Jie","Li Lei"))
    mutate(golden,record,args.outdir/"missing-sources",lambda prs:replace_all(prs,"Source:","Source omitted:"))
    image=args.outdir/"flat.png"; Image.new("RGB",(1600,900),"white").save(image)
    def flattened(prs):
        slide=prs.slides[2]
        for shape in list(slide.shapes): shape._element.getparent().remove(shape._element)
        slide.shapes.add_picture(str(image),0,0,width=prs.slide_width,height=prs.slide_height)
    mutate(golden,record,args.outdir/"flattened-software-slide",flattened)
    def marker(prs):
        box=prs.slides[-1].shapes.add_textbox(Inches(.1),Inches(.1),Inches(2),Inches(.2)); box.text="solution/Golden"
    mutate(golden,record,args.outdir/"hardcoded-marker",marker)
    output=directory(args.outdir,"extra-output"); shutil.copy2(golden,output/PPTX_NAME); shutil.copy2(record,output/JSON_NAME); (output/"debug.txt").write_text("debug\n")
if __name__=="__main__": main()
