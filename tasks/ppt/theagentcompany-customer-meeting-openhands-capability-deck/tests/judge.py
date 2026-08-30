#!/usr/bin/env python3
"""TheAgentCompany OpenHands 客户会前演示任务的确定性 Judge。"""
from __future__ import annotations

import argparse
import hashlib
import json
import posixpath
import re
import subprocess
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

PPTX_NAME="openhands_intro.pptx"; JSON_NAME="Delivery_Record.json"
TASK_ID="cowork-evo/theagentcompany-customer-meeting-openhands-capability-deck"
RUBRIC_VERSION="theagentcompany-openhands-customer-deck-adaptation-v0.1"; JUDGE_VERSION="theagentcompany-openhands-customer-deck-deterministic-v0.1"
HASHES={"Original_Request.md":"eb8a006788c291de4c76a47e8f72a33670395d1f2999c99b45c7e0dd8c5bbdef","RocketChat_Recipient_Context.json":"865c354cfc784ef51994a0b0159dd396f3a76c3f8ba01046a039d3d0d1ac07b8","OpenHands_Evaluation_Snapshot.json":"aa411f23fbb791720b0c9274a5900685c13d929f2dce1c20d711a406a26d3338"}
EXPECTED_DELIVERY={"schema_version":"0.1.0","artifact_filename":PPTX_NAME,"original_workspace_path":"/workspace/openhands_intro.pptx","recipient":"Huang Jie","channel":"RocketChat direct message","message":PPTX_NAME,"request_purpose":"customer meeting materials","source_reference":"http://the-agent-company.com:8929/root/openhands/-/tree/main/evaluation?ref_type=heads","network_side_effect":"not_executed_offline"}
SUSPICIOUS=["solution/golden","judge-result","artifact_dir","file:///root","/root/solution","expected_delivery","upstream-evaluator.py"]
SPECS={"R001":(6,"hurdle"),"R002":(10,"hurdle"),"R003":(14,"hurdle"),"R004":(12,"hurdle"),"R005":(14,"hurdle"),"R006":(10,"hurdle"),"R007":(12,"hurdle"),"R008":(20,"hurdle"),"R009":(2,"positive"),"R010":(-20,"penalty")}

def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def norm(value): return re.sub(r"\s+"," ",str(value).replace("\u00a0"," ").replace("\v"," ")).strip()
def locate(name):
    for path in (Path("/app/input_files")/name,Path(__file__).resolve().parents[1]/"data/input_files"/name):
        if path.is_file(): return path
    return None
def source_ok():
    missing=[]
    for name,digest in HASHES.items():
        path=locate(name)
        if path is None or sha(path)!=digest: missing.append(name)
    snapshot=None
    if not missing:
        snapshot=json.loads(locate("OpenHands_Evaluation_Snapshot.json").read_text()); categories=snapshot.get("categories",[])
        if [item.get("order") for item in categories]!=[1,2,3] or [len(item.get("benchmarks",[])) for item in categories]!=[8,2,6] or snapshot.get("derived_totals")!={"domains":3,"benchmarks":16,"software_engineering":8,"web_browsing":2,"misc_assistance":6}: missing.append("snapshot-shape")
        recipient=json.loads(locate("RocketChat_Recipient_Context.json").read_text())
        if "Huang Jie" not in recipient: missing.append("recipient-context")
    return not missing,missing,snapshot
def package(path):
    evidence={"broken":None,"missing":[],"macros":[],"external":[]}
    try:
        with zipfile.ZipFile(path) as archive:
            evidence["broken"]=archive.testzip(); names=set(archive.namelist()); evidence["macros"]=[name for name in names if name.lower().endswith("vbaproject.bin")]
            for relation in sorted(name for name in names if name.endswith(".rels")):
                root=etree.fromstring(archive.read(relation)); owner="" if relation=="_rels/.rels" else posixpath.join(relation.split("/_rels/",1)[0],relation.split("/_rels/",1)[1][:-5]); base=posixpath.dirname(owner)
                for rel in root:
                    target=rel.attrib.get("Target","")
                    if rel.attrib.get("TargetMode")=="External": evidence["external"].append(target); continue
                    resolved=target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base,target))
                    if resolved not in names: evidence["missing"].append(f"{relation}->{resolved}")
        return evidence["broken"] is None and not evidence["missing"] and not evidence["macros"],evidence
    except Exception as exc: evidence["exception"]=str(exc); return False,evidence
def slide_text(slide): return norm("\n".join(shape.text for shape in slide.shapes if getattr(shape,"has_text_frame",False)))
def missing(text,items): return [value for value in items if value.lower() not in text.lower()]
def render(path,out):
    out.mkdir(parents=True,exist_ok=True); profile=out/"profile"; profile.mkdir(exist_ok=True); process=subprocess.run(["soffice",f"-env:UserInstallation=file://{profile}","--headless","--convert-to","pdf","--outdir",str(out),str(path)],stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,timeout=300); pdf=out/f"{path.stem}.pdf"
    if process.returncode or not pdf.is_file(): return 0,process.stdout[-500:]
    info=subprocess.run(["pdfinfo",str(pdf)],stdout=subprocess.PIPE,text=True).stdout; match=re.search(r"^Pages:\s+(\d+)",info,re.M); return int(match.group(1)) if match else 0,process.stdout[-500:]
def item(cid,passed,evidence):
    weight,kind=SPECS[cid]; return {"criterion_id":cid,"score":1.0 if passed else 0.0,"evidence":evidence,"method":"deterministic","raw":{"weight":weight,"type":kind,"passed":bool(passed)}}
def assemble(paths,criteria):
    hurdles=[entry["criterion_id"] for entry in criteria if entry["raw"]["type"]=="hurdle" and not entry["raw"]["passed"]]; total=sum(entry["raw"]["weight"] for entry in criteria if entry["raw"]["weight"]>0); earned=sum(entry["raw"]["weight"] for entry in criteria if entry["raw"]["weight"]>0 and entry["raw"]["passed"]); penalty=sum(abs(entry["raw"]["weight"]) for entry in criteria if entry["raw"]["weight"]<0 and not entry["raw"]["passed"]); reward=0 if hurdles else max(0,min(1,(earned-penalty)/total)); return {"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":[sha(path) for path in paths if path.is_file()],"criterion_results":criteria,"reward":reward,"passed":reward>=.8 and not hurdles,"hurdle_failures":hurdles}
def evaluate(pptx_path,json_path,artifact):
    src,source_evidence,snapshot=source_ok()
    try: prs=Presentation(pptx_path)
    except Exception as exc: return assemble([pptx_path,json_path],[item(cid,False,f"打开失败:{exc}") for cid in SPECS])
    try: delivery=json.loads(json_path.read_text())
    except Exception as exc: delivery={"parse_error":str(exc)}
    pkg,pkg_evidence=package(pptx_path); ratio=prs.slide_width/prs.slide_height if prs.slide_height else 0; slides=[slide_text(slide) for slide in prs.slides]; alltext="\n".join(slides)
    r1=src and pkg and json_path.is_file() and len(prs.slides)==6 and abs(ratio-16/9)<.02; criteria=[item("R001",r1,f"pages={len(prs.slides)} ratio={ratio} source={source_evidence} package={json.dumps(pkg_evidence,ensure_ascii=False)} delivery_parse={delivery.get('parse_error')}")]
    if snapshot is None: snapshot={"categories":[]}
    cover_missing=missing(slides[0] if slides else "",["OpenHands Agent","Capability Inventory","3","task domains","16","benchmark entries","0","performance claims","frozen GitLab evaluation tree"])
    map_missing=missing(slides[1] if len(slides)>1 else "",["Software Engineering","8","Web Browsing","2","Misc. Assistance","6","Customer-safe statement"])
    r2=not cover_missing and not map_missing; criteria.append(item("R002",r2,f"cover_missing={cover_missing} map_missing={map_missing}"))
    categories=snapshot.get("categories",[]); all_names=[name for category in categories for name in category.get("benchmarks",[])]
    domain_results=[]
    for offset,category in enumerate(categories,2):
        page=slides[offset] if offset<len(slides) else ""; own=category["benchmarks"]; others=[name for name in all_names if name not in own]; absent=missing(page,[category["name"],category["upstream_evaluator_anchor"],*own]); foreign=[name for name in others if name.lower() in page.lower()]; domain_results.append((not absent and not foreign,absent,foreign))
    while len(domain_results)<3: domain_results.append((False,["missing-domain-page"],[]))
    criteria.append(item("R003",domain_results[0][0],f"missing={domain_results[0][1]} foreign={domain_results[0][2]}")); criteria.append(item("R004",domain_results[1][0],f"missing={domain_results[1][1]} foreign={domain_results[1][2]}")); criteria.append(item("R005",domain_results[2][0],f"missing={domain_results[2][1]} foreign={domain_results[2][2]}"))
    sequence=[("Software Engineering","software engineer","8 benchmark entries"),("Web Browsing","web browsing","2 benchmark entries"),("Misc. Assistance","misc assistance","6 benchmark entries")]; sequence_errors=[]
    for index,(title,anchor,count) in enumerate(sequence,2):
        page=slides[index] if index<len(slides) else ""
        for value in (title,anchor,count):
            if value.lower() not in page.lower(): sequence_errors.append(f"page={index+1}:{value}")
    r6=not sequence_errors; criteria.append(item("R006",r6,f"errors={sequence_errors}"))
    boundary_required=["Inventory coverage is not measured performance","No score, success rate, model version, or comparative ranking","What it does not establish: pass rate, model version, score, ranking, or customer outcome","Names + domain membership only. No scores, success rates, versions, rankings, or customer outcomes"]
    boundary_missing=missing(alltext,boundary_required); unsupported=re.findall(r"\b(?:achieves?|achieved|outperforms?|state-of-the-art|accuracy\s+of|success\s+rate\s+of|pass\s+rate\s+of)\b|\b\d+(?:\.\d+)?%",alltext,re.I)
    r7=not boundary_missing and not unsupported; criteria.append(item("R007",r7,f"missing={boundary_missing} unsupported={unsupported}"))
    delivery_ok=delivery==EXPECTED_DELIVERY; last=slides[5] if len(slides)>5 else ""; handoff_missing=missing(last,["Evidence & Handoff",PPTX_NAME,"Huang Jie","RocketChat DM","No live network side effect is claimed","not_executed_offline",EXPECTED_DELIVERY["source_reference"]]); citations=sum(text.count("Source:") for text in slides); tags=sum("98b68ef8" in text for text in slides); markers=[value for value in SUSPICIOUS if value in alltext.lower()]; pictures=[shape.shape_id for slide in prs.slides for shape in slide.shapes if shape.shape_type==MSO_SHAPE_TYPE.PICTURE]; native_text=sum(getattr(shape,"has_text_frame",False) for slide in prs.slides for shape in slide.shapes)
    with zipfile.ZipFile(pptx_path) as archive: media=[name for name in archive.namelist() if name.startswith("ppt/media/") and not name.endswith("/")]
    r8=delivery_ok and not handoff_missing and citations==6 and tags==6 and not markers and native_text>=80 and not pictures and not media and not pkg_evidence.get("macros") and not pkg_evidence.get("external"); criteria.append(item("R008",r8,f"delivery_ok={delivery_ok} handoff_missing={handoff_missing} citations={citations} tags={tags} markers={markers} native_text={native_text} pictures={pictures} media={media}"))
    pages=0; log=""
    try:
        with tempfile.TemporaryDirectory(prefix="tac-openhands-render-") as directory: pages,log=render(pptx_path,Path(directory))
    except Exception as exc: log=str(exc)
    delivered=sorted(str(path.relative_to(artifact)) for path in artifact.rglob("*") if path.is_file()); r9=pages==6 and delivered==[JSON_NAME,PPTX_NAME] and pkg; criteria.append(item("R009",r9,f"pages={pages} files={delivered} log={log}"))
    severe=not all([r2,*[result[0] for result in domain_results],r6,r7,r8]); criteria.append(item("R010",not severe,f"severe={severe}")); return assemble([pptx_path,json_path],criteria)
def main():
    parser=argparse.ArgumentParser(); parser.add_argument("--artifact-dir",type=Path,required=True); parser.add_argument("--result",type=Path,required=True); parser.add_argument("--reward-file",type=Path,required=True); args=parser.parse_args(); pptx=args.artifact_dir/PPTX_NAME; record=args.artifact_dir/JSON_NAME
    result=evaluate(pptx,record,args.artifact_dir) if pptx.is_file() and record.is_file() else assemble([pptx,record],[item(cid,False,f"缺失交付文件：{PPTX_NAME} / {JSON_NAME}") for cid in SPECS]); args.result.parent.mkdir(parents=True,exist_ok=True); args.reward_file.parent.mkdir(parents=True,exist_ok=True); args.result.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n"); args.reward_file.write_text(f"{result['reward']:.6f}\n"); print(json.dumps(result,ensure_ascii=False,indent=2))
if __name__=="__main__": main()
