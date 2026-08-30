#!/usr/bin/env python3
"""TheAgentCompany GitLab engineering progress 任务的逐题 Golden。"""
from __future__ import annotations
import argparse,hashlib,json
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.util import Inches,Pt

HASHES={"GitLab_Repository_Snapshot.json":"b48461179a363e8b70cb92d3692ed7dc53ae67cbfe76984627dc04c8c597a729","Original_Request.md":"c8f04427aa726c69a25ab593c0ef65ea6a0a9503d79c9138d6b4afa296d57a37"}
NAVY=RGBColor(0x11,0x1D,0x32); BLUE=RGBColor(0x2E,0x68,0xC5); CYAN=RGBColor(0x18,0xA5,0xB7); GREEN=RGBColor(0x18,0x8B,0x64); AMBER=RGBColor(0xDB,0x88,0x16); RED=RGBColor(0xC7,0x3B,0x35); WHITE=RGBColor(0xFF,0xFF,0xFF); INK=RGBColor(0x20,0x2A,0x38); MUTED=RGBColor(0x68,0x74,0x86); LINE=RGBColor(0xDF,0xE5,0xEC); SURFACE=RGBColor(0xF4,0xF7,0xFA); PALE=RGBColor(0xE9,0xF1,0xFB)
def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def text(slide,value,left,top,width,height,*,size=14,color=INK,bold=False,align=PP_ALIGN.LEFT,margin=.04,valign=MSO_ANCHOR.TOP):
    shape=slide.shapes.add_textbox(Inches(left),Inches(top),Inches(width),Inches(height)); frame=shape.text_frame; frame.clear(); frame.word_wrap=True; frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=Inches(margin); frame.vertical_anchor=valign; paragraph=frame.paragraphs[0]; paragraph.alignment=align; run=paragraph.add_run(); run.text=str(value); run.font.name="Aptos"; run.font.size=Pt(size); run.font.color.rgb=color; run.font.bold=bold; return shape
def rect(slide,left,top,width,height,fill,line=None,rounded=True):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(left),Inches(top),Inches(width),Inches(height)); shape.fill.solid(); shape.fill.fore_color.rgb=fill; shape.line.color.rgb=line or fill; return shape
def bg(slide,color): slide.background.fill.solid(); slide.background.fill.fore_color.rgb=color
def footer(slide,page,dark=False):
    color=RGBColor(0xA8,0xB5,0xC8) if dark else MUTED; text(slide,"Source: GitLab_Repository_Snapshot.json  ·  upstream evaluator @ 98b68ef8",.65,7.12,10.9,.18,size=7.2,color=color); text(slide,page,12.15,7.12,.45,.18,size=8,color=color,align=PP_ALIGN.RIGHT)
def metric(slide,value,label,left,top,width,accent):
    rect(slide,left,top,width,1.36,WHITE,LINE); rect(slide,left,top,.07,1.36,accent,rounded=False); text(slide,f"{value:,}",left+.24,top+.18,width-.42,.47,size=28,color=NAVY,bold=True); text(slide,label.upper(),left+.24,top+.82,width-.42,.23,size=8.5,color=MUTED,bold=True); text(slide,"ALL STATES",left+.24,top+1.08,width-.42,.17,size=7.3,color=accent,bold=True)
def verify(root):
    for name,digest in HASHES.items():
        path=root/name
        if not path.is_file() or sha(path)!=digest: raise ValueError(f"冻结输入错误：{name}")
    snapshot=json.loads((root/"GitLab_Repository_Snapshot.json").read_text()); repos=snapshot["repositories"]
    if [repo["order"] for repo in repos]!=list(range(1,15)) or sum(repo["issues"] for repo in repos)!=4990 or sum(repo["merge_requests"] for repo in repos)!=794: raise ValueError("仓库快照不一致")
    return snapshot
def build(input_dir,output):
    snapshot=verify(input_dir); repos=snapshot["repositories"]; prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]
    slide=prs.slides.add_slide(blank); bg(slide,NAVY)
    text(slide,"CTO BRIEFING  /  FROZEN GITLAB SNAPSHOT",.72,.61,7.3,.22,size=9.5,color=CYAN,bold=True); text(slide,"Overview of Our\nEngineering Progress",.72,1.14,7.8,1.42,size=35,color=WHITE,bold=True); text(slide,"A complete repository-by-repository activity register.",.75,2.93,6.9,.32,size=14,color=RGBColor(0xC7,0xD2,0xE1)); text(slide,snapshot["gitlab_url"],.75,3.44,5.6,.26,size=10.5,color=CYAN,bold=True)
    text(slide,"Prepared for Sarah Johnson, CTO",.75,3.82,5.6,.22,size=9.5,color=RGBColor(0xB9,0xC5,0xD7),bold=True)
    values=[("14","repositories",BLUE),("4,990","issues",GREEN),("794","merge requests",AMBER),("10","repos with activity",CYAN)]
    for i,(value,label,accent) in enumerate(values):
        left=.72+i*3.04; rect(slide,left,4.55,2.78,1.35,RGBColor(0x1B,0x2D,0x4D),RGBColor(0x31,0x48,0x70)); rect(slide,left,4.55,.06,1.35,accent,rounded=False); text(slide,value,left+.22,4.75,2.28,.38,size=23,color=WHITE,bold=True); text(slide,label.upper(),left+.22,5.35,2.28,.2,size=8,color=accent,bold=True)
    text(slide,"Counts include open + closed issues and open + merged/closed merge requests. They describe activity volume, not quality or current backlog.",.75,6.32,11.4,.36,size=9,color=RGBColor(0xB9,0xC5,0xD7)); footer(slide,1,True)

    max_issues=max(repo["issues"] for repo in repos); max_mr=max(repo["merge_requests"] for repo in repos)
    for index,repo in enumerate(repos,1):
        slide=prs.slides.add_slide(blank); bg(slide,SURFACE); accent=[BLUE,CYAN,GREEN,AMBER][(index-1)%4]
        text(slide,f"REPOSITORY {index:02d} / 14",.68,.3,3.2,.2,size=8.5,color=accent,bold=True); text(slide,repo["name"],.68,.62,9.4,.58,size=29,color=NAVY,bold=True); text(slide,f"{snapshot['gitlab_url']}root/{repo['name']}",.7,1.3,7.0,.24,size=9.5,color=BLUE,bold=True); text(slide,f"{index+1:02d} / 15",11.7,.36,.92,.2,size=8.5,color=MUTED,align=PP_ALIGN.RIGHT); rect(slide,.68,1.76,12.0,.02,LINE,rounded=False)
        description=repo["description"] or "No description set in frozen GitLab snapshot"
        rect(slide,.72,2.12,7.0,2.03,WHITE,LINE); text(slide,"DESCRIPTION",.98,2.4,1.25,.2,size=8.5,color=accent,bold=True); text(slide,description,.98,2.88,6.45,.84,size=17 if len(description)<70 else 14,color=NAVY,bold=True)
        if not repo["description"]:
            rect(slide,.98,3.78,2.08,.25,RGBColor(0xFB,0xEC,0xE9),RGBColor(0xFB,0xEC,0xE9)); text(slide,"METADATA GAP · NOT INFERRED",1.06,3.82,1.92,.14,size=6.8,color=RED,bold=True,align=PP_ALIGN.CENTER)
        metric(slide,repo["issues"],"Issues",8.02,2.12,2.13,GREEN); metric(slide,repo["merge_requests"],"Merge Requests",10.42,2.12,2.13,AMBER)
        text(slide,"ACTIVITY FOOTPRINT",.75,4.63,2.0,.2,size=8.5,color=MUTED,bold=True)
        labels=[("Issues",repo["issues"],max_issues,GREEN),("Merge requests",repo["merge_requests"],max_mr,AMBER)]
        for row,(label,value,maximum,color) in enumerate(labels):
            top=5.05+row*.62; text(slide,label,.75,top,1.45,.2,size=9,color=INK,bold=True); rect(slide,2.28,top+.02,8.75,.16,LINE,LINE,rounded=False); width=8.75*(value/maximum) if maximum else 0
            if width>0: rect(slide,2.28,top+.02,max(.05,width),.16,color,color,rounded=False)
            text(slide,f"{value:,}",11.28,top-.03,1.0,.24,size=10,color=NAVY,bold=True,align=PP_ALIGN.RIGHT)
        rect(slide,.72,6.36,11.83,.4,PALE,PALE); text(slide,"Frozen exact count  ·  no +/-5 live-image tolerance  ·  repository order preserved",.95,6.46,11.4,.18,size=8.2,color=BLUE,bold=True,align=PP_ALIGN.CENTER); footer(slide,index+1)
    output.parent.mkdir(parents=True,exist_ok=True); prs.save(output)
    if len(Presentation(output).slides)!=15: raise ValueError("页数不匹配")
    print(f"{output}\nsha256={sha(output)}")
def main():
    parser=argparse.ArgumentParser(); parser.add_argument("--input-dir",type=Path,required=True); parser.add_argument("--output",type=Path,required=True); args=parser.parse_args(); build(args.input_dir,args.output)
if __name__=="__main__": main()
