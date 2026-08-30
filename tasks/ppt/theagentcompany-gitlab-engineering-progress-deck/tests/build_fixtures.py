#!/usr/bin/env python3
from __future__ import annotations

import argparse
import shutil
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

PPTX_NAME="Engineering_Progress.pptx"

def directory(root,name):
    path=root/name; path.mkdir(parents=True,exist_ok=True); return path

def replace_on_slide(slide,old,new):
    count=0
    for shape in slide.shapes:
        if not getattr(shape,"has_text_frame",False): continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if old in run.text: run.text=run.text.replace(old,new); count+=1
    if not count: raise ValueError(f"未找到需要变异的文本：{old}")

def replace_all(prs,old,new):
    count=0
    for slide in prs.slides:
        for shape in slide.shapes:
            if not getattr(shape,"has_text_frame",False): continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if old in run.text: run.text=run.text.replace(old,new); count+=1
    if not count: raise ValueError(f"未找到需要变异的文本：{old}")

def mutate(golden,target,operation):
    target=directory(target.parent,target.name); prs=Presentation(golden); operation(prs); prs.save(target/PPTX_NAME)

def swap_slides(prs,first,second):
    ids=prs.slides._sldIdLst
    a,b=ids[first],ids[second]
    ids.remove(a); ids.remove(b)
    ids.insert(first,b); ids.insert(second,a)

def delete_slide(prs,index):
    ids=prs.slides._sldIdLst; slide_id=ids[index]; prs.part.drop_rel(slide_id.rId); ids.remove(slide_id)

def main():
    parser=argparse.ArgumentParser(); parser.add_argument("--golden-dir",type=Path,required=True); parser.add_argument("--outdir",type=Path,required=True); args=parser.parse_args(); golden=args.golden_dir/PPTX_NAME
    directory(args.outdir,"empty")
    output=directory(args.outdir,"corrupt"); (output/PPTX_NAME).write_bytes(b"bad\n")
    output=directory(args.outdir,"valid-repackage")
    with zipfile.ZipFile(golden) as source,zipfile.ZipFile(output/PPTX_NAME,"w",compression=zipfile.ZIP_STORED) as target:
        for info in source.infolist(): target.writestr(info.filename,source.read(info.filename))
    def alt_cover(prs): prs.slides[0].background.fill.solid(); prs.slides[0].background.fill.fore_color.rgb=RGBColor(0x16,0x28,0x3A)
    mutate(golden,args.outdir/"valid-alt-cover",alt_cover)
    mutate(golden,args.outdir/"swapped-repo-slides",lambda prs:swap_slides(prs,1,2))
    mutate(golden,args.outdir/"missing-repo-slide",lambda prs:delete_slide(prs,4))
    mutate(golden,args.outdir/"wrong-description",lambda prs:replace_on_slide(prs.slides[6],"JanusGraph: an open-source, distributed graph database","A cloud-native graph analytics platform"))
    mutate(golden,args.outdir/"invented-empty-description",lambda prs:replace_on_slide(prs.slides[1],"No description set in frozen GitLab snapshot","Internal API gateway for all company products"))
    mutate(golden,args.outdir/"wrong-issue-by-one",lambda prs:replace_on_slide(prs.slides[2],"28","29"))
    mutate(golden,args.outdir/"wrong-mr-by-one",lambda prs:replace_on_slide(prs.slides[7],"277","278"))
    mutate(golden,args.outdir/"wrong-cover-total",lambda prs:replace_on_slide(prs.slides[0],"4,990","4,991"))
    mutate(golden,args.outdir/"wrong-count-scope",lambda prs:replace_on_slide(prs.slides[0],"open + closed issues and open + merged/closed merge requests","open issues and open merge requests"))
    mutate(golden,args.outdir/"repo-open-only",lambda prs:replace_on_slide(prs.slides[8],"ALL STATES","OPEN ONLY"))
    mutate(golden,args.outdir/"missing-sources",lambda prs:replace_all(prs,"Source:","Source omitted:"))
    image=args.outdir/"flat.png"; Image.new("RGB",(1600,900),"white").save(image)
    def flattened(prs):
        slide=prs.slides[10]
        for shape in list(slide.shapes): shape._element.getparent().remove(shape._element)
        slide.shapes.add_picture(str(image),0,0,width=prs.slide_width,height=prs.slide_height)
    mutate(golden,args.outdir/"flattened-repo-slide",flattened)
    def marker(prs):
        box=prs.slides[-1].shapes.add_textbox(Inches(.1),Inches(.1),Inches(2),Inches(.2)); box.text="solution/Golden"
    mutate(golden,args.outdir/"hardcoded-marker",marker)
    output=directory(args.outdir,"extra-output"); shutil.copy2(golden,output/PPTX_NAME); (output/"debug.txt").write_text("debug\n")

if __name__=="__main__": main()
