#!/usr/bin/env python3
"""TheAgentCompany GitLab 工程进展任务的逐页确定性 Judge。"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import posixpath
import re
import subprocess
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

PPTX_NAME="Engineering_Progress.pptx"
TASK_ID="cowork-evo/theagentcompany-gitlab-engineering-progress-deck"
RUBRIC_VERSION="theagentcompany-gitlab-progress-ppt-adaptation-v0.1"
JUDGE_VERSION="theagentcompany-gitlab-progress-deterministic-v0.1"
HASHES={"Original_Request.md":"c8f04427aa726c69a25ab593c0ef65ea6a0a9503d79c9138d6b4afa296d57a37","GitLab_Repository_Snapshot.json":"b48461179a363e8b70cb92d3692ed7dc53ae67cbfe76984627dc04c8c597a729"}
SUSPICIOUS=["solution/golden","judge-result","artifact_dir","file:///root","/root/solution","expected_repos","upstream-evaluator.py"]
SPECS={"R001":(6,"hurdle"),"R002":(8,"hurdle"),"R003":(12,"hurdle"),"R004":(14,"hurdle"),"R005":(18,"hurdle"),"R006":(18,"hurdle"),"R007":(12,"hurdle"),"R008":(10,"hurdle"),"R009":(2,"positive"),"R010":(-20,"penalty")}

def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def norm(value): return re.sub(r"\s+"," ",str(value).replace("\u00a0"," ").replace("\v"," ")).strip()

def locate(name):
    for path in (Path("/app/input_files")/name,Path(__file__).resolve().parents[1]/"data/input_files"/name):
        if path.is_file(): return path
    return None

def source_ok():
    missing=[]
    for name,digest in HASHES.items():
        path=locate(name)
        if path is None or sha(path)!=digest: missing.append(name)
    snapshot=None
    if not missing:
        snapshot=json.loads(locate("GitLab_Repository_Snapshot.json").read_text())
        repos=snapshot.get("repositories",[])
        if [repo.get("order") for repo in repos]!=list(range(1,15)): missing.append("repository-order")
        if snapshot.get("derived_totals")!={"repositories":14,"repositories_with_activity":10,"issues":4990,"merge_requests":794}: missing.append("derived-totals")
        if sum(repo.get("issues",-1) for repo in repos)!=4990 or sum(repo.get("merge_requests",-1) for repo in repos)!=794: missing.append("repository-totals")
    return not missing,missing,snapshot

def package(path):
    evidence={"broken":None,"missing":[],"macros":[],"external":[]}
    try:
        with zipfile.ZipFile(path) as archive:
            evidence["broken"]=archive.testzip(); names=set(archive.namelist()); evidence["macros"]=[n for n in names if n.lower().endswith("vbaproject.bin")]
            for relation in sorted(n for n in names if n.endswith(".rels")):
                root=etree.fromstring(archive.read(relation)); owner="" if relation=="_rels/.rels" else posixpath.join(relation.split("/_rels/",1)[0],relation.split("/_rels/",1)[1][:-5]); base=posixpath.dirname(owner)
                for rel in root:
                    target=rel.attrib.get("Target","")
                    if rel.attrib.get("TargetMode")=="External": evidence["external"].append(target); continue
                    resolved=target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base,target))
                    if resolved not in names: evidence["missing"].append(f"{relation}->{resolved}")
        return evidence["broken"] is None and not evidence["missing"] and not evidence["macros"],evidence
    except Exception as exc:
        evidence["exception"]=str(exc); return False,evidence

def slide_text(slide): return norm("\n".join(shape.text for shape in slide.shapes if getattr(shape,"has_text_frame",False)))

def has_repo_name(text,name):
    return re.search(rf"(?<![A-Za-z0-9_.-]){re.escape(name)}(?![A-Za-z0-9_.-])",text,re.I) is not None

def number_from_text(value):
    match=re.search(r"(?<![A-Za-z0-9])([0-9]{1,3}(?:,[0-9]{3})*|[0-9]+)(?![A-Za-z0-9])",value)
    return int(match.group(1).replace(",","")) if match else None

def metric_value(slide,label):
    """先支持“Issues: 28”同框写法，再通过位置匹配分开的标签和数字。"""
    label_lower=label.lower(); shapes=[]
    for shape in slide.shapes:
        if not getattr(shape,"has_text_frame",False): continue
        value=norm(shape.text); shapes.append((shape,value))
        if label_lower in value.lower():
            patterns=[rf"{re.escape(label)}\s*:?\s*([0-9][0-9,]*)",rf"([0-9][0-9,]*)\s*{re.escape(label)}"]
            for pattern in patterns:
                match=re.search(pattern,value,re.I)
                if match: return int(match.group(1).replace(",",""))
    candidates=[]
    for label_shape,label_text in shapes:
        if label_lower not in label_text.lower(): continue
        lx=label_shape.left+label_shape.width/2; ly=label_shape.top+label_shape.height/2
        for shape,value in shapes:
            if shape is label_shape or not re.fullmatch(r"[0-9][0-9,]*",value): continue
            dx=abs((shape.left+shape.width/2)-lx); dy=abs((shape.top+shape.height/2)-ly)
            if dx<=4*914400 and dy<=1.6*914400:
                candidates.append((math.hypot(dx,dy),int(value.replace(",",""))))
    return min(candidates)[1] if candidates else None

def render(path,out):
    out.mkdir(parents=True,exist_ok=True); profile=out/"profile"; profile.mkdir(exist_ok=True)
    process=subprocess.run(["soffice",f"-env:UserInstallation=file://{profile}","--headless","--convert-to","pdf","--outdir",str(out),str(path)],stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,timeout=300); pdf=out/f"{path.stem}.pdf"
    if process.returncode or not pdf.is_file(): return 0,process.stdout[-500:]
    info=subprocess.run(["pdfinfo",str(pdf)],stdout=subprocess.PIPE,text=True).stdout; match=re.search(r"^Pages:\s+(\d+)",info,re.M); return int(match.group(1)) if match else 0,process.stdout[-500:]

def item(cid,passed,evidence):
    weight,kind=SPECS[cid]; return {"criterion_id":cid,"score":1.0 if passed else 0.0,"evidence":evidence,"method":"deterministic","raw":{"weight":weight,"type":kind,"passed":bool(passed)}}

def assemble(path,criteria):
    hurdles=[entry["criterion_id"] for entry in criteria if entry["raw"]["type"]=="hurdle" and not entry["raw"]["passed"]]
    total=sum(entry["raw"]["weight"] for entry in criteria if entry["raw"]["weight"]>0); earned=sum(entry["raw"]["weight"] for entry in criteria if entry["raw"]["weight"]>0 and entry["raw"]["passed"]); penalty=sum(abs(entry["raw"]["weight"]) for entry in criteria if entry["raw"]["weight"]<0 and not entry["raw"]["passed"])
    reward=0 if hurdles else max(0,min(1,(earned-penalty)/total))
    return {"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":[sha(path)] if path.is_file() else [],"criterion_results":criteria,"reward":reward,"passed":reward>=.8 and not hurdles,"hurdle_failures":hurdles}

def evaluate(path,artifact):
    src,source_evidence,snapshot=source_ok()
    try: prs=Presentation(path)
    except Exception as exc: return assemble(path,[item(cid,False,f"打开失败:{exc}") for cid in SPECS])
    pkg,pkg_evidence=package(path); ratio=prs.slide_width/prs.slide_height if prs.slide_height else 0
    r1=src and pkg and len(prs.slides)==15 and abs(ratio-16/9)<.02
    criteria=[item("R001",r1,f"pages={len(prs.slides)} ratio={ratio} source={source_evidence} package={json.dumps(pkg_evidence,ensure_ascii=False)}")]
    if snapshot is None: snapshot={"repositories":[],"derived_totals":{}}
    slides=[slide_text(slide) for slide in prs.slides]
    cover=slides[0] if slides else ""
    cover_required=["Overview of Our Engineering Progress","Sarah Johnson","14","repositories","4,990","issues","794","merge requests","10","repos with activity"]
    missing_cover=[value for value in cover_required if value.lower() not in cover.lower()]
    r2=not missing_cover; criteria.append(item("R002",r2,f"missing={missing_cover}"))

    order_errors=[]; description_errors=[]; issue_errors=[]; mr_errors=[]
    repos=snapshot.get("repositories",[])
    for index,repo in enumerate(repos,1):
        if index>=len(slides):
            order_errors.append(f"{index}:{repo['name']}:missing-slide"); description_errors.append(repo["name"]); issue_errors.append(f"{repo['name']}:missing"); mr_errors.append(f"{repo['name']}:missing"); continue
        page=slides[index]
        found=[candidate["name"] for candidate in repos if has_repo_name(page,candidate["name"])]
        if not has_repo_name(page,repo["name"]) or found!=[repo["name"]]: order_errors.append(f"{index}:{repo['name']}:found={found}")
        expected_description=repo["description"] or "No description set in frozen GitLab snapshot"
        if norm(expected_description).lower() not in page.lower(): description_errors.append(repo["name"])
        actual_issues=metric_value(prs.slides[index],"Issues")
        actual_mr=metric_value(prs.slides[index],"Merge Requests")
        if actual_issues!=repo["issues"]: issue_errors.append(f"{repo['name']}:{actual_issues}!={repo['issues']}")
        if actual_mr!=repo["merge_requests"]: mr_errors.append(f"{repo['name']}:{actual_mr}!={repo['merge_requests']}")
    r3=len(repos)==14 and not order_errors; criteria.append(item("R003",r3,f"errors={order_errors}"))
    r4=len(repos)==14 and not description_errors; criteria.append(item("R004",r4,f"errors={description_errors}"))
    r5=len(repos)==14 and not issue_errors; criteria.append(item("R005",r5,f"errors={issue_errors}"))
    r6=len(repos)==14 and not mr_errors; criteria.append(item("R006",r6,f"errors={mr_errors}"))

    scope_required=["open + closed issues","open + merged/closed merge requests","activity volume","not quality","not quality or current backlog","no +/-5 live-image tolerance"]
    scope_text="\n".join(slides)
    scope_hits={value:value.lower() in scope_text.lower() for value in scope_required}
    per_page_scope=[index+1 for index,text in enumerate(slides[1:]) if text.lower().count("all states")<2]
    r7=all(scope_hits.values()) and not per_page_scope; criteria.append(item("R007",r7,f"scope={scope_hits} pages_missing_two_all_states={per_page_scope}"))

    alltext="\n".join(slides); citations=sum(text.count("Source:") for text in slides); commit_tags=sum("98b68ef8" in text for text in slides); markers=[value for value in SUSPICIOUS if value in alltext.lower()]
    pictures=[shape.shape_id for slide in prs.slides for shape in slide.shapes if shape.shape_type==MSO_SHAPE_TYPE.PICTURE]; native_text=sum(getattr(shape,"has_text_frame",False) for slide in prs.slides for shape in slide.shapes)
    with zipfile.ZipFile(path) as archive: media=[name for name in archive.namelist() if name.startswith("ppt/media/") and not name.endswith("/")]
    r8=citations==15 and commit_tags==15 and not markers and native_text>=100 and not pictures and not media and not pkg_evidence.get("macros") and not pkg_evidence.get("external")
    criteria.append(item("R008",r8,f"citations={citations} commit_tags={commit_tags} markers={markers} native_text={native_text} pictures={pictures} media={media}"))

    pages=0; log=""
    try:
        with tempfile.TemporaryDirectory(prefix="tac-gitlab-render-") as directory: pages,log=render(path,Path(directory))
    except Exception as exc: log=str(exc)
    delivered=sorted(str(file.relative_to(artifact)) for file in artifact.rglob("*") if file.is_file())
    r9=pages==15 and delivered==[PPTX_NAME] and pkg; criteria.append(item("R009",r9,f"pages={pages} files={delivered} log={log}"))
    severe=not all([r2,r3,r4,r5,r6,r7,r8]); criteria.append(item("R010",not severe,f"severe={severe}"))
    return assemble(path,criteria)

def main():
    parser=argparse.ArgumentParser(); parser.add_argument("--artifact-dir",type=Path,required=True); parser.add_argument("--result",type=Path,required=True); parser.add_argument("--reward-file",type=Path,required=True); args=parser.parse_args(); path=args.artifact_dir/PPTX_NAME
    result=evaluate(path,args.artifact_dir) if path.is_file() else assemble(path,[item(cid,False,f"缺失交付文件：{PPTX_NAME}") for cid in SPECS])
    args.result.parent.mkdir(parents=True,exist_ok=True); args.reward_file.parent.mkdir(parents=True,exist_ok=True); args.result.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n"); args.reward_file.write_text(f"{result['reward']:.6f}\n"); print(json.dumps(result,ensure_ascii=False,indent=2))

if __name__=="__main__": main()
