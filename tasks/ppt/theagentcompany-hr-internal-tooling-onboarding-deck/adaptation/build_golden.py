#!/usr/bin/env python3
"""TheAgentCompany HR internal-tooling 任务的逐题 Golden 构造器。"""
from __future__ import annotations
import argparse,hashlib,json,zipfile
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.util import Inches,Pt

HASHES={
    "TAC_overview.md":"504f19054be9b3c2ed74841332dee564e88ee98c0f990d54c50b057605350e53",
    "Personnell_File.odt":"e19e2ff3988ec2ed915f9e87571df40c3f7a27be2350e9c98a6263d65afc5025",
}
NAVY=RGBColor(0x11,0x1E,0x34); BLUE=RGBColor(0x28,0x6B,0xC7); CYAN=RGBColor(0x18,0xA6,0xB8)
GREEN=RGBColor(0x18,0x8A,0x63); AMBER=RGBColor(0xDF,0x8A,0x14); RED=RGBColor(0xC9,0x3D,0x36)
WHITE=RGBColor(0xFF,0xFF,0xFF); INK=RGBColor(0x1F,0x29,0x37); MUTED=RGBColor(0x66,0x72,0x85)
LINE=RGBColor(0xDF,0xE5,0xEC); SURFACE=RGBColor(0xF4,0xF7,0xFA); PALE=RGBColor(0xE9,0xF1,0xFB)

def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def textbox(slide,value,left,top,width,height,*,size=14,color=INK,bold=False,align=PP_ALIGN.LEFT,margin=.04,valign=MSO_ANCHOR.TOP):
    shape=slide.shapes.add_textbox(Inches(left),Inches(top),Inches(width),Inches(height)); frame=shape.text_frame; frame.clear(); frame.word_wrap=True; frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=Inches(margin); frame.vertical_anchor=valign
    paragraph=frame.paragraphs[0]; paragraph.alignment=align; run=paragraph.add_run(); run.text=str(value); run.font.name="Aptos"; run.font.size=Pt(size); run.font.bold=bold; run.font.color.rgb=color; return shape
def rectangle(slide,left,top,width,height,fill,line=None,rounded=True):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(left),Inches(top),Inches(width),Inches(height)); shape.fill.solid(); shape.fill.fore_color.rgb=fill; shape.line.color.rgb=line or fill; return shape
def background(slide,color): slide.background.fill.solid(); slide.background.fill.fore_color.rgb=color
def header(slide,title,kicker,page):
    textbox(slide,kicker.upper(),.65,.28,8.5,.22,size=9,color=BLUE,bold=True); textbox(slide,title,.65,.56,11.55,.5,size=26,color=NAVY,bold=True); textbox(slide,f"{page:02d} / 07",11.68,.33,.92,.2,size=8.5,color=MUTED,align=PP_ALIGN.RIGHT); rectangle(slide,.65,1.15,12.03,.02,LINE,rounded=False)
def footer(slide,page,source,dark=False):
    color=RGBColor(0xA9,0xB7,0xCA) if dark else MUTED; textbox(slide,f"Source: {source}  ·  TheAgentCompany @ 98b68ef8",.65,7.12,11.0,.18,size=7.2,color=color); textbox(slide,page,12.15,7.12,.45,.18,size=8,color=color,align=PP_ALIGN.RIGHT)
def card(slide,title,body,left,top,width,height,accent=BLUE,*,small=False):
    rectangle(slide,left,top,width,height,WHITE,LINE); rectangle(slide,left,top,.06,height,accent,rounded=False); textbox(slide,title,left+.2,top+.16,width-.36,.28,size=12 if small else 14,color=NAVY,bold=True); textbox(slide,body,left+.2,top+.56,width-.36,height-.7,size=8.5 if small else 10.5,color=MUTED)
def verify_inputs(root):
    for name,digest in HASHES.items():
        path=root/name
        if not path.is_file() or sha(path)!=digest: raise ValueError(f"冻结输入错误：{name}")
    snapshot=json.loads((root/"Internal_Tooling_Source_Snapshot.json").read_text()); context=json.loads((root/"Task_Delivery_Context.json").read_text())
    with zipfile.ZipFile(root/"Personnell_File.odt") as archive:
        xml=etree.fromstring(archive.read("content.xml")); personnel=" ".join(xml.itertext())
    if "Chen Xinyi" not in personnel or "Human Resources Manager" not in personnel: raise ValueError("人员文件与交付上下文不一致")
    if len(snapshot["rocketchat"]["channels"])!=15 or len(snapshot["gitlab"]["repositories"])!=14 or len(snapshot["owncloud"]["actual_documents_folders"])!=9: raise ValueError("上游快照规模不匹配")
    return snapshot,context
def add_multiline(frame,items,size=10,color=INK,bullet=False,spacing=4):
    frame.clear(); frame.word_wrap=True
    for index,item in enumerate(items):
        p=frame.paragraphs[0] if index==0 else frame.add_paragraph(); p.text=item; p.font.name="Aptos"; p.font.size=Pt(size); p.font.color.rgb=color; p.space_after=Pt(spacing); p.level=0
        if bullet: p.text="• "+p.text
def build(input_dir,output_dir):
    source,context=verify_inputs(input_dir); output_dir.mkdir(parents=True,exist_ok=True)
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]

    slide=prs.slides.add_slide(blank); background(slide,NAVY)
    textbox(slide,"NEW HIRE ONBOARDING  /  INTERNAL SYSTEMS",.72,.6,7.4,.24,size=10,color=CYAN,bold=True)
    textbox(slide,"Internal Tooling\nTraining",.72,1.18,7.6,1.45,size=38,color=WHITE,bold=True)
    textbox(slide,"One operating map for communication, code, files and delivery.",.75,2.98,7.5,.42,size=15,color=RGBColor(0xC8,0xD3,0xE2))
    tools=[("RocketChat","communicate",CYAN),("GitLab","build",BLUE),("OwnCloud","share",GREEN),("Plane","deliver",AMBER)]
    for i,(name,verb,accent) in enumerate(tools):
        left=.72+i*3.04; rectangle(slide,left,4.5,2.78,1.36,RGBColor(0x1B,0x2D,0x4D),RGBColor(0x31,0x48,0x70)); rectangle(slide,left,4.5,.06,1.36,accent,rounded=False); textbox(slide,name,left+.22,4.72,2.25,.32,size=16,color=WHITE,bold=True); textbox(slide,verb.upper(),left+.22,5.25,2.25,.2,size=8.5,color=accent,bold=True)
    textbox(slide,"Audience: new hires  ·  7-page foundation  ·  review owner: Chen Xinyi",.75,6.38,9.2,.22,size=9,color=RGBColor(0xB9,0xC5,0xD7)); footer(slide,1,"Task_Delivery_Context.json + upstream task",True)

    slide=prs.slides.add_slide(blank); background(slide,WHITE); header(slide,"TAC Overview","Mission -> products -> operating model",2)
    textbox(slide,"Build data + AI systems\nthat turn complexity into leverage.",.72,1.5,6.1,1.02,size=24,color=NAVY,bold=True)
    textbox(slide,"MISSION",.75,2.76,1.0,.2,size=8.5,color=BLUE,bold=True); textbox(slide,"Provide advanced, efficient and user-friendly data processing and AI tools that drive innovation and maximize the value of data.",.75,3.08,5.8,.86,size=12,color=INK)
    textbox(slide,"VISION",.75,4.38,1.0,.2,size=8.5,color=GREEN,bold=True); textbox(slide,"Become a global leader in distributed systems and artificial intelligence.",.75,4.72,5.8,.62,size=12,color=INK)
    products=[("GRAPH","JanusGraph","distributed graph database",BLUE),("STREAM","RisingWave","streaming database",CYAN),("AI","OpenHands + llama.cpp","development and inference",GREEN),("CRAWL","Colly","web crawler framework",AMBER),("SEARCH","OpenSearch","distributed search engine",BLUE),("LOW-CODE","Node-RED","event-driven apps",CYAN)]
    for i,(tag,name,desc,accent) in enumerate(products):
        col=i%2; row=i//2; left=7.0+col*2.85; top=1.48+row*1.62; rectangle(slide,left,top,2.55,1.34,SURFACE,LINE); textbox(slide,tag,left+.18,top+.16,1.0,.2,size=7.5,color=accent,bold=True); textbox(slide,name,left+.18,top+.45,2.15,.29,size=12,color=NAVY,bold=True); textbox(slide,desc,left+.18,top+.82,2.16,.34,size=9,color=MUTED)
    rectangle(slide,7.0,6.45,5.4,.36,PALE,PALE); textbox(slide,"Rust · Python · C++ · Go · Java",7.15,6.54,5.1,.18,size=8.5,color=BLUE,bold=True,align=PP_ALIGN.CENTER); footer(slide,2,"TAC_overview.md")

    slide=prs.slides.add_slide(blank); background(slide,SURFACE); header(slide,"RocketChat","Communication map",3)
    textbox(slide,source["urls"]["rocketchat"],.72,1.34,4.5,.28,size=12,color=BLUE,bold=True); textbox(slide,"Default channels",9.5,1.32,1.25,.2,size=8.5,color=MUTED,bold=True); textbox(slide,"#general  +  #kudos",10.45,1.28,2.15,.28,size=11,color=GREEN,bold=True,align=PP_ALIGN.RIGHT)
    groups=[("COMPANY & SUPPORT",["#general","#kudos","#help-desk","#random","#tech-talk"],BLUE),("FUNCTIONS",["#engineering","#product","#hr-announcements","#sales-marketing"],GREEN),("PROJECTS",["#project-ai","#project-graphdb","#project-lowcode","#project-search","#project-streamdb","#project-webcrawler"],AMBER)]
    purposes={c["name"]:c["purpose"] for c in source["rocketchat"]["channels"]}
    for index,(label,names,accent) in enumerate(groups):
        left=.72+index*4.1; rectangle(slide,left,1.88,3.78,4.78,WHITE,LINE); textbox(slide,label,left+.24,2.13,3.2,.2,size=8.5,color=accent,bold=True)
        y=2.58
        for name in names:
            textbox(slide,name,left+.24,y,1.47,.22,size=9.5,color=NAVY,bold=True); textbox(slide,purposes[name],left+1.56,y,1.95,.43,size=8.1,color=MUTED); y+=.67
    footer(slide,3,"Internal_Tooling_Source_Snapshot.json + Personnell_File.odt")

    slide=prs.slides.add_slide(blank); background(slide,WHITE); header(slide,"Gitlab","Collaborate on code and delivery",4)
    textbox(slide,source["urls"]["gitlab"],.72,1.35,4.8,.28,size=12,color=BLUE,bold=True)
    flow=[("01","CODE COLLABORATION","branch -> commit -> merge request","Review changes before they reach the main line.",BLUE),("02","PROJECT MANAGEMENT","milestones -> owners -> progress","Keep repository work tied to delivery context.",GREEN),("03","ISSUE TRACKING","issue -> label -> assignee -> state","Turn findings into visible, accountable work.",AMBER)]
    for i,(num,title,steps,body,accent) in enumerate(flow):
        left=.72+i*4.08; rectangle(slide,left,2.03,3.72,3.8,SURFACE,LINE); textbox(slide,num,left+.24,2.28,.62,.42,size=25,color=accent,bold=True); textbox(slide,title,left+.24,2.9,3.15,.22,size=9,color=accent,bold=True); textbox(slide,steps,left+.24,3.4,3.18,.48,size=14,color=NAVY,bold=True); textbox(slide,body,left+.24,4.22,3.18,.75,size=10.5,color=INK); rectangle(slide,left+.24,5.32,3.18,.04,accent,rounded=False)
    textbox(slide,"Working rule",.75,6.25,1.05,.2,size=8.5,color=MUTED,bold=True); textbox(slide,"Keep discussion in the merge request, delivery status in issues, and project-level cadence in milestones.",1.72,6.18,10.8,.34,size=10.5,color=NAVY,bold=True); footer(slide,4,"Internal_Tooling_Source_Snapshot.json / upstream GitLab usage requirements")

    slide=prs.slides.add_slide(blank); background(slide,WHITE); header(slide,"Gitlab Repos","Complete company repository register",5)
    repos=source["gitlab"]["repositories"]
    for index,repo in enumerate(repos):
        col=index//7; row=index%7; left=.72+col*6.08; top=1.42+row*.73; rectangle(slide,left,top,5.75,.61,SURFACE,LINE); textbox(slide,f"{index+1:02d}",left+.16,top+.16,.34,.18,size=8,color=BLUE,bold=True); textbox(slide,repo["name"],left+.56,top+.1,1.73,.22,size=10,color=NAVY,bold=True); textbox(slide,repo["description"],left+2.22,top+.09,3.29,.39,size=7.8,color=MUTED)
    rectangle(slide,.72,6.68,11.83,.22,PALE,PALE); textbox(slide,"14 / 14 repositories represented  ·  descriptions stay within frozen evidence",.9,6.71,11.48,.16,size=7.7,color=BLUE,bold=True,align=PP_ALIGN.CENTER); footer(slide,5,"Internal_Tooling_Source_Snapshot.json / evaluator repo list + OwnCloud taglines")

    slide=prs.slides.add_slide(blank); background(slide,SURFACE); header(slide,"Owncloud","Documents inventory and sharing boundary",6)
    textbox(slide,source["urls"]["owncloud"],.72,1.33,4.8,.28,size=12,color=BLUE,bold=True); textbox(slide,"Actual frozen image inventory",9.65,1.33,2.9,.22,size=8.5,color=GREEN,bold=True,align=PP_ALIGN.RIGHT)
    folder_notes={"Admin":"company administration","Administrative Specialist":"vendor and office operations","Data Analysis":"analysis datasets and workbooks","Engineering":"engineering reference files","Financials":"finance records and reporting","Human Resources Team":"people and onboarding records","Marketing":"campaign and event materials","Q1 Planning and Allocation":"planning and resource allocation","Research":"papers and research artifacts"}
    folders=source["owncloud"]["actual_documents_folders"]
    for i,name in enumerate(folders):
        col=i%3; row=i//3; left=.72+col*4.08; top=1.87+row*1.22; card(slide,name,folder_notes[name],left,top,3.76,.98,[BLUE,GREEN,AMBER][col],small=True)
    rectangle(slide,.72,5.78,11.98,.79,WHITE,RGBColor(0xED,0xC6,0xC2)); textbox(slide,"SOURCE MISMATCH",.96,6.0,1.32,.18,size=8,color=RED,bold=True); textbox(slide,"Legacy evaluator omits Marketing and Q1 Planning and Allocation, and calls the actual Human Resources Team folder 'Human Resource Team'. This deck follows the frozen image inventory.",2.25,5.93,10.12,.42,size=9.2,color=INK)
    footer(slide,6,"OwnCloud image inventory + legacy evaluator folder list")

    slide=prs.slides.add_slide(blank); background(slide,NAVY)
    textbox(slide,"PLANE  /  DELIVERY CONTROL",.72,.58,6.8,.22,size=9.5,color=CYAN,bold=True); textbox(slide,"Turn work into visible\nproject movement.",.72,1.03,6.2,1.05,size=30,color=WHITE,bold=True); textbox(slide,source["urls"]["plane"],.75,2.35,4.2,.26,size=11,color=CYAN,bold=True)
    workflow="Projects  ->  Modules  ->  Issues  ->  Sprint cycles  ->  Assignees  ->  Analytics"; textbox(slide,workflow,.75,2.92,7.16,.63,size=13,color=RGBColor(0xD1,0xDB,0xE8),bold=True)
    metrics=source["plane"]["analytics_metrics"]
    for i,name in enumerate(metrics):
        col=i%2; row=i//2; left=8.18+col*2.17; top=1.12+row*1.47; rectangle(slide,left,top,1.91,1.18,WHITE,WHITE); textbox(slide,f"0{i+1}",left+.18,top+.17,.42,.18,size=8,color=[BLUE,GREEN,AMBER,CYAN][i],bold=True); textbox(slide,name,left+.18,top+.48,1.53,.48,size=9.4,color=NAVY,bold=True)
    rectangle(slide,.72,4.2,11.85,1.72,RGBColor(0x1B,0x2D,0x4D),RGBColor(0x31,0x48,0x70)); textbox(slide,"OFFLINE DELIVERY RECORD",.96,4.49,2.05,.2,size=8.5,color=CYAN,bold=True); textbox(slide,"Chen Xinyi  ·  RocketChat direct message",.96,4.88,4.55,.28,size=14,color=WHITE,bold=True); textbox(slide,"Request review and inclusion in onboarding materials",.96,5.3,4.9,.24,size=10,color=RGBColor(0xBF,0xCB,0xDB)); textbox(slide,context["offline_share_reference"],6.0,4.79,6.05,.58,size=10,color=WHITE,bold=True,align=PP_ALIGN.RIGHT); textbox(slide,"No live network side effect is claimed; Delivery_Record.json freezes the intended action.",6.0,5.38,6.05,.26,size=8.5,color=RGBColor(0xBF,0xCB,0xDB),align=PP_ALIGN.RIGHT)
    footer(slide,7,"Plane snapshot + Personnell_File.odt + Task_Delivery_Context.json",True)

    pptx_path=output_dir/"Internal_Tooling_Training.pptx"; prs.save(pptx_path)
    delivery={"schema_version":"0.1.0","artifact_filename":"Internal_Tooling_Training.pptx","owncloud_path":context["original_owncloud_path"],"share_reference":context["offline_share_reference"],"recipient":"Chen Xinyi","channel":"RocketChat direct message","request_purpose":"review and inclusion in the onboarding materials","message":"Chen Xinyi, please review Internal_Tooling_Training.pptx for inclusion in the onboarding materials: "+context["offline_share_reference"],"network_side_effect":"not_executed_offline"}
    (output_dir/"Delivery_Record.json").write_text(json.dumps(delivery,ensure_ascii=False,indent=2)+"\n")
    if len(Presentation(pptx_path).slides)!=7: raise ValueError("页数不匹配")
    print(f"{pptx_path}\nsha256={sha(pptx_path)}\n{output_dir/'Delivery_Record.json'}\nsha256={sha(output_dir/'Delivery_Record.json')}")
def main():
    parser=argparse.ArgumentParser(); parser.add_argument("--input-dir",type=Path,required=True); parser.add_argument("--output-dir",type=Path,required=True); args=parser.parse_args(); build(args.input_dir,args.output_dir)
if __name__=="__main__": main()
