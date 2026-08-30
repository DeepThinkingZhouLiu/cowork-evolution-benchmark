#!/usr/bin/env python3
from __future__ import annotations
import argparse,json,shutil,zipfile
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

PPTX_NAME="Internal_Tooling_Training.pptx"; JSON_NAME="Delivery_Record.json"
def directory(root,name): path=root/name; path.mkdir(parents=True,exist_ok=True); return path
def replace(prs,old,new):
    count=0
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape,"has_text_frame",False):
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if old in run.text: run.text=run.text.replace(old,new); count+=1
    if not count: raise ValueError(old)
def mutate(golden,target,operation,record):
    target=directory(target.parent,target.name); prs=Presentation(golden); operation(prs); prs.save(target/PPTX_NAME); shutil.copy2(record,target/JSON_NAME)
def mutate_json(golden,record,target,key,value):
    target=directory(target.parent,target.name); shutil.copy2(golden,target/PPTX_NAME); data=json.loads(record.read_text()); data[key]=value; (target/JSON_NAME).write_text(json.dumps(data,ensure_ascii=False,indent=2)+"\n")
def main():
    parser=argparse.ArgumentParser(); parser.add_argument("--golden-dir",type=Path,required=True); parser.add_argument("--outdir",type=Path,required=True); args=parser.parse_args(); golden=args.golden_dir/PPTX_NAME; record=args.golden_dir/JSON_NAME
    directory(args.outdir,"empty"); output=directory(args.outdir,"corrupt"); (output/PPTX_NAME).write_bytes(b"bad\n"); shutil.copy2(record,output/JSON_NAME)
    output=directory(args.outdir,"ppt-only"); shutil.copy2(golden,output/PPTX_NAME)
    output=directory(args.outdir,"json-only"); shutil.copy2(record,output/JSON_NAME)
    output=directory(args.outdir,"valid-repackage")
    with zipfile.ZipFile(golden) as source,zipfile.ZipFile(output/PPTX_NAME,"w",compression=zipfile.ZIP_STORED) as target:
        for info in source.infolist(): target.writestr(info.filename,source.read(info.filename))
    shutil.copy2(record,output/JSON_NAME)
    def alt_cover(prs): prs.slides[0].background.fill.solid(); prs.slides[0].background.fill.fore_color.rgb=RGBColor(0x0C,0x32,0x3D)
    mutate(golden,args.outdir/"valid-alt-cover",alt_cover,record)
    mutate(golden,args.outdir/"wrong-tac-product",lambda prs:replace(prs,"JanusGraph","Cassandra"),record)
    mutate(golden,args.outdir/"missing-channel",lambda prs:replace(prs,"#project-webcrawler","#channel-omitted"),record)
    mutate(golden,args.outdir/"wrong-default-channel",lambda prs:replace(prs,"#general  +  #kudos","#general  +  #random"),record)
    mutate(golden,args.outdir/"missing-repo",lambda prs:replace(prs,"copilot-arena-server","repo omitted"),record)
    mutate(golden,args.outdir/"wrong-folder",lambda prs:replace(prs,"Q1 Planning and Allocation","Q2 Forecasts"),record)
    mutate(golden,args.outdir/"missing-evaluator-caveat",lambda prs:replace(prs,"Legacy evaluator omits Marketing and Q1 Planning and Allocation","Folder list verified"),record)
    mutate(golden,args.outdir/"wrong-plane-metric",lambda prs:replace(prs,"Most Issues Closed","Average Cycle Time"),record)
    mutate_json(golden,record,args.outdir/"wrong-recipient","recipient","Sarah Johnson")
    mutate_json(golden,record,args.outdir/"wrong-share-path","share_reference","owncloud:/Documents/Internal_Tooling_Training.pptx")
    mutate(golden,args.outdir/"cross-artifact-mismatch",lambda prs:replace(prs,"owncloud:/Documents/Human Resources Team/Internal_Tooling_Training.pptx","owncloud:/Documents/Training/Internal_Tooling_Training.pptx"),record)
    mutate(golden,args.outdir/"missing-citations",lambda prs:replace(prs,"Source:","Sources omitted:"),record)
    image=args.outdir/"flat.png"; Image.new("RGB",(1600,900),"white").save(image)
    def flattened(prs):
        slide=prs.slides[2]
        for shape in list(slide.shapes): shape._element.getparent().remove(shape._element)
        slide.shapes.add_picture(str(image),0,0,width=prs.slide_width,height=prs.slide_height)
    mutate(golden,args.outdir/"flattened-rocketchat",flattened,record)
    def marker(prs): box=prs.slides[-1].shapes.add_textbox(Inches(.1),Inches(.1),Inches(2),Inches(.2)); box.text="solution/Golden"
    mutate(golden,args.outdir/"hardcoded-marker",marker,record)
    output=directory(args.outdir,"extra-output"); shutil.copy2(golden,output/PPTX_NAME); shutil.copy2(record,output/JSON_NAME); (output/"debug.txt").write_text("debug\n")
if __name__=="__main__": main()
