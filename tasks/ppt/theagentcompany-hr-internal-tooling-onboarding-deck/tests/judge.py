#!/usr/bin/env python3
"""TheAgentCompany internal-tooling 任务专属确定性 Judge。"""
from __future__ import annotations
import argparse,hashlib,json,posixpath,re,subprocess,tempfile,zipfile
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

PPTX_NAME="Internal_Tooling_Training.pptx"; JSON_NAME="Delivery_Record.json"
TASK_ID="cowork-evo/theagentcompany-hr-internal-tooling-onboarding-deck"; RUBRIC_VERSION="theagentcompany-internal-tooling-ppt-adaptation-v0.1"; JUDGE_VERSION="theagentcompany-internal-tooling-deterministic-v0.1"
HASHES={"TAC_overview.md":"504f19054be9b3c2ed74841332dee564e88ee98c0f990d54c50b057605350e53","Personnell_File.odt":"e19e2ff3988ec2ed915f9e87571df40c3f7a27be2350e9c98a6263d65afc5025","Internal_Tooling_Source_Snapshot.json":"f6d8baf345c00fb7c7e72277688ce2ecc706c390498b18025428624ca25eae2e","Task_Delivery_Context.json":"cb2725244476dac2ec672c7c60585c408f664ffc5937dc19b47e1bff712d8311"}
PRODUCTS=["JanusGraph","RisingWave","OpenHands + llama.cpp","Colly","OpenSearch","Node-RED"]
CHANNELS=["#general","#kudos","#engineering","#help-desk","#hr-announcements","#product","#project-ai","#project-graphdb","#project-lowcode","#project-search","#project-streamdb","#project-webcrawler","#random","#sales-marketing","#tech-talk"]
REPOS=["api-server","bustub","colly","copilot-arena-server","Documentation","janusgraph","llama.cpp","node-red","openhands","opensearch","raft","risingwave","sotopia","streamlit"]
FOLDERS=["Admin","Administrative Specialist","Data Analysis","Engineering","Financials","Human Resources Team","Marketing","Q1 Planning and Allocation","Research"]
PLANE=["Projects","Modules","Issues","Sprint cycles","Assignees","Analytics","Total Open Tasks","Pending Issues","Most Issues Created","Most Issues Closed"]
EXPECTED_DELIVERY={"schema_version":"0.1.0","artifact_filename":PPTX_NAME,"owncloud_path":"/Documents/Human Resources Team/Internal_Tooling_Training.pptx","share_reference":"owncloud:/Documents/Human Resources Team/Internal_Tooling_Training.pptx","recipient":"Chen Xinyi","channel":"RocketChat direct message","request_purpose":"review and inclusion in the onboarding materials","message":"Chen Xinyi, please review Internal_Tooling_Training.pptx for inclusion in the onboarding materials: owncloud:/Documents/Human Resources Team/Internal_Tooling_Training.pptx","network_side_effect":"not_executed_offline"}
SUSPICIOUS=["solution/golden","judge-result","artifact_dir","/root/","expected_delivery"]
SPECS={"R001":(6,"hurdle"),"R002":(12,"hurdle"),"R003":(14,"hurdle"),"R004":(10,"hurdle"),"R005":(14,"hurdle"),"R006":(14,"hurdle"),"R007":(12,"hurdle"),"R008":(16,"hurdle"),"R009":(2,"positive"),"R010":(-20,"penalty")}

def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def norm(value): return re.sub(r"\s+"," ",str(value).replace("\u00a0"," ").replace("\v"," ")).strip()
def locate(name):
    for path in (Path("/app/input_files")/name,Path(__file__).resolve().parents[1]/"data/input_files"/name):
        if path.is_file(): return path
    return None
def source_ok():
    missing=[]
    for name,digest in HASHES.items():
        path=locate(name)
        if path is None or sha(path)!=digest: missing.append(name)
    if not missing:
        snapshot=json.loads(locate("Internal_Tooling_Source_Snapshot.json").read_text()); context=json.loads(locate("Task_Delivery_Context.json").read_text())
        if len(snapshot["rocketchat"]["channels"])!=15 or len(snapshot["gitlab"]["repositories"])!=14 or len(snapshot["owncloud"]["actual_documents_folders"])!=9: missing.append("snapshot-shape")
        if context.get("hr_manager")!="Chen Xinyi": missing.append("delivery-context")
    return not missing,missing
def package(path):
    evidence={"broken":None,"missing":[],"macros":[],"external":[]}
    try:
        with zipfile.ZipFile(path) as archive:
            evidence["broken"]=archive.testzip(); names=set(archive.namelist()); evidence["macros"]=[n for n in names if n.lower().endswith("vbaproject.bin")]
            for relation in sorted(n for n in names if n.endswith(".rels")):
                root=etree.fromstring(archive.read(relation)); owner="" if relation=="_rels/.rels" else posixpath.join(relation.split("/_rels/",1)[0],relation.split("/_rels/",1)[1][:-5]); base=posixpath.dirname(owner)
                for rel in root:
                    target=rel.attrib.get("Target","")
                    if rel.attrib.get("TargetMode")=="External": evidence["external"].append(target); continue
                    resolved=target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base,target))
                    if resolved not in names: evidence["missing"].append(f"{relation}->{resolved}")
        return evidence["broken"] is None and not evidence["missing"] and not evidence["macros"],evidence
    except Exception as exc: evidence["exception"]=str(exc); return False,evidence
def slide_text(slide): return norm("\n".join(shape.text for shape in slide.shapes if getattr(shape,"has_text_frame",False)))
def all_text(prs): return "\n".join(slide_text(slide) for slide in prs.slides)
def render(path,out):
    out.mkdir(parents=True,exist_ok=True); profile=out/"profile"; profile.mkdir(exist_ok=True); process=subprocess.run(["soffice",f"-env:UserInstallation=file://{profile}","--headless","--convert-to","pdf","--outdir",str(out),str(path)],stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,timeout=300); pdf=out/f"{path.stem}.pdf"
    if process.returncode or not pdf.is_file(): return 0,process.stdout[-500:]
    info=subprocess.run(["pdfinfo",str(pdf)],stdout=subprocess.PIPE,text=True).stdout; match=re.search(r"^Pages:\s+(\d+)",info,re.M); return int(match.group(1)) if match else 0,process.stdout[-500:]
def item(cid,passed,evidence):
    weight,kind=SPECS[cid]; return {"criterion_id":cid,"score":1.0 if passed else 0.0,"evidence":evidence,"method":"deterministic","raw":{"weight":weight,"type":kind,"passed":bool(passed)}}
def assemble(path,criteria):
    hurdles=[x["criterion_id"] for x in criteria if x["raw"]["type"]=="hurdle" and not x["raw"]["passed"]]; total=sum(x["raw"]["weight"] for x in criteria if x["raw"]["weight"]>0); earned=sum(x["raw"]["weight"] for x in criteria if x["raw"]["weight"]>0 and x["raw"]["passed"]); penalty=sum(abs(x["raw"]["weight"]) for x in criteria if x["raw"]["weight"]<0 and not x["raw"]["passed"]); reward=0 if hurdles else max(0,min(1,(earned-penalty)/total)); return {"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":[sha(path)] if path.is_file() else [],"criterion_results":criteria,"reward":reward,"passed":reward>=.8 and not hurdles,"hurdle_failures":hurdles}
def contains_all(text,items): return [value for value in items if value.lower() not in text.lower()]
def evaluate(pptx_path,json_path,artifact):
    src,source_evidence=source_ok()
    try: prs=Presentation(pptx_path)
    except Exception as exc: return assemble(pptx_path,[item(cid,False,f"打开失败:{exc}") for cid in SPECS])
    try: delivery=json.loads(json_path.read_text())
    except Exception as exc: delivery={"parse_error":str(exc)}
    pkg,pkg_evidence=package(pptx_path); ratio=prs.slide_width/prs.slide_height if prs.slide_height else 0; r1=src and pkg and json_path.is_file() and len(prs.slides)==7 and abs(ratio-16/9)<.02; criteria=[item("R001",r1,f"pages={len(prs.slides)} ratio={ratio} source={source_evidence} package={json.dumps(pkg_evidence,ensure_ascii=False)} delivery_parse={delivery.get('parse_error')}")]
    slides=[slide_text(slide) for slide in prs.slides]; text="\n".join(slides)
    missing_overview=contains_all(slides[1],["TAC Overview","MISSION","VISION","Provide advanced, efficient and user-friendly data processing and AI tools","global leader in distributed systems and artificial intelligence",*PRODUCTS]); r2=not missing_overview; criteria.append(item("R002",r2,f"missing={missing_overview}"))
    missing_channels=contains_all(slides[2],["RocketChat","http://the-agent-company.com:3000/","Default channels","#general + #kudos",*CHANNELS]); purposes=["company-wide communication","recognition and team achievements","engineering coordination","internal support requests","HR announcements","product planning and UX coordination","project work","informal team conversation","sales and marketing coordination","cross-project technical discussion"]; missing_purposes=contains_all(slides[2],purposes); r3=not missing_channels and not missing_purposes; criteria.append(item("R003",r3,f"missing_channels={missing_channels} missing_purposes={missing_purposes}"))
    missing_gitlab=contains_all(slides[3],["Gitlab","http://the-agent-company.com:8929/","code collaboration","project management","issue tracking","branch -> commit -> merge request","issue -> label -> assignee -> state"]); r4=not missing_gitlab; criteria.append(item("R004",r4,f"missing={missing_gitlab}"))
    missing_repos=contains_all(slides[4],["Gitlab Repos",*REPOS]); unexpected=[name for name in ["Cassandra","MongoDB","Jira","Bitbucket"] if name.lower() in slides[4].lower()]; r5=not missing_repos and not unexpected; criteria.append(item("R005",r5,f"missing={missing_repos} unexpected={unexpected}"))
    missing_folders=contains_all(slides[5],["Owncloud","http://the-agent-company.com:8092/",*FOLDERS,"Legacy evaluator omits Marketing and Q1 Planning and Allocation","Human Resource Team","frozen image inventory"]); r6=not missing_folders; criteria.append(item("R006",r6,f"missing={missing_folders}"))
    missing_plane=contains_all(slides[6],["Plane","http://the-agent-company.com:8091/",*PLANE]); r7=not missing_plane; criteria.append(item("R007",r7,f"missing={missing_plane}"))
    delivery_ok=delivery==EXPECTED_DELIVERY and EXPECTED_DELIVERY["share_reference"] in slides[6] and "Chen Xinyi" in slides[6] and "No live network side effect is claimed" in slides[6]; citations=text.count("Source:"); tags=text.count("TheAgentCompany @ 98b68ef8"); markers=[value for value in SUSPICIOUS if value in text.lower()]; pictures=[shape.shape_id for slide in prs.slides for shape in slide.shapes if shape.shape_type==MSO_SHAPE_TYPE.PICTURE]; native_text=sum(getattr(shape,"has_text_frame",False) for slide in prs.slides for shape in slide.shapes)
    with zipfile.ZipFile(pptx_path) as archive: media=[name for name in archive.namelist() if name.startswith("ppt/media/") and not name.endswith("/")]
    r8=delivery_ok and citations==7 and tags==7 and not markers and native_text>=70 and not pictures and not media and not pkg_evidence.get("macros") and not pkg_evidence.get("external"); criteria.append(item("R008",r8,f"delivery_ok={delivery_ok} citations={citations} tags={tags} markers={markers} native_text={native_text} pictures={pictures} media={media}"))
    pages=0; log=""
    try:
        with tempfile.TemporaryDirectory(prefix="tac-tooling-render-") as directory: pages,log=render(pptx_path,Path(directory))
    except Exception as exc: log=str(exc)
    delivered=sorted(str(path.relative_to(artifact)) for path in artifact.rglob("*") if path.is_file()); r9=pages==7 and delivered==[JSON_NAME,PPTX_NAME] and pkg; criteria.append(item("R009",r9,f"pages={pages} files={delivered} log={log}"))
    severe=not all([r2,r3,r4,r5,r6,r7,r8]); criteria.append(item("R010",not severe,f"severe={severe}")); return assemble(pptx_path,criteria)
def main():
    parser=argparse.ArgumentParser(); parser.add_argument("--artifact-dir",type=Path,required=True); parser.add_argument("--result",type=Path,required=True); parser.add_argument("--reward-file",type=Path,required=True); args=parser.parse_args(); pptx=args.artifact_dir/PPTX_NAME; record=args.artifact_dir/JSON_NAME
    result=evaluate(pptx,record,args.artifact_dir) if pptx.is_file() and record.is_file() else assemble(pptx,[item(cid,False,f"缺失交付文件：{PPTX_NAME} / {JSON_NAME}") for cid in SPECS]); args.result.parent.mkdir(parents=True,exist_ok=True); args.reward_file.parent.mkdir(parents=True,exist_ok=True); args.result.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n"); args.reward_file.write_text(f"{result['reward']:.6f}\n"); print(json.dumps(result,ensure_ascii=False,indent=2))
if __name__=="__main__": main()
