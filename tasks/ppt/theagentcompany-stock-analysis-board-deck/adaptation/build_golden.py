#!/usr/bin/env python3
"""TheAgentCompany 董事会股价分析题的逐题 Golden。"""
from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

HASHES={"Original_Request.md":"621281d20fa3e0b4dbf6903ff871e035c99bb194a80b4fe3e1188f7f7954281b","stock_data_raw.csv":"5e83e1084cdebcfe78892cc345d38a8d584d2533e58efab722294a94edf1874c"}
CLEAN_HASH="1b8940d298205bed6700b1f7aea0321385c5aff0599f3fcec173300c9bb3fcbf"
NAVY=RGBColor(0x0E,0x1D,0x2D); BLUE=RGBColor(0x2F,0x6F,0xD4); CYAN=RGBColor(0x13,0xA7,0xB8); GREEN=RGBColor(0x18,0x8E,0x68); AMBER=RGBColor(0xDF,0x8B,0x16); RED=RGBColor(0xC7,0x43,0x3B); WHITE=RGBColor(0xFF,0xFF,0xFF); INK=RGBColor(0x17,0x22,0x34); MUTED=RGBColor(0x66,0x73,0x85); LINE=RGBColor(0xDE,0xE5,0xEC); SURFACE=RGBColor(0xF4,0xF7,0xFA); PALE=RGBColor(0xE8,0xF1,0xFB)

def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def text(slide,value,left,top,width,height,*,size=14,color=INK,bold=False,align=PP_ALIGN.LEFT,margin=.04,valign=MSO_ANCHOR.TOP):
    shape=slide.shapes.add_textbox(Inches(left),Inches(top),Inches(width),Inches(height)); frame=shape.text_frame; frame.clear(); frame.word_wrap=True; frame.auto_size=MSO_AUTO_SIZE.NONE; frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=Inches(margin); frame.vertical_anchor=valign; paragraph=frame.paragraphs[0]; paragraph.alignment=align; run=paragraph.add_run(); run.text=str(value); run.font.name="Aptos"; run.font.size=Pt(size); run.font.color.rgb=color; run.font.bold=bold; return shape
def rect(slide,left,top,width,height,fill,line=None,rounded=True):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(left),Inches(top),Inches(width),Inches(height)); shape.fill.solid(); shape.fill.fore_color.rgb=fill; shape.line.color.rgb=line or fill; return shape
def bg(slide): slide.background.fill.solid(); slide.background.fill.fore_color.rgb=SURFACE
def footer(slide,page): text(slide,"Source: stock_cleaned.csv  ·  derived from stock_data_raw.csv @ OwnCloud 1.0.0",.65,7.12,11.3,.18,size=7.2,color=MUTED); text(slide,str(page),12.2,7.12,.35,.18,size=8,color=MUTED,align=PP_ALIGN.RIGHT)
def header(slide,kicker,title,page): bg(slide); text(slide,kicker,.68,.34,7.2,.24,size=8.2,color=BLUE,bold=True); text(slide,title,.68,.72,11.8,.52,size=24 if len(title)>42 else 27,color=NAVY,bold=True); rect(slide,.68,1.43,12,.02,LINE,rounded=False); footer(slide,page)
def note_card(slide,label,value,top,accent=BLUE):
    rect(slide,9.42,top,3.18,.9,WHITE,LINE); rect(slide,9.42,top,.06,.9,accent,accent,rounded=False); text(slide,label.upper(),9.68,top+.14,2.5,.16,size=7.2,color=MUTED,bold=True); text(slide,value,9.68,top+.41,2.55,.34,size=12 if len(str(value))>20 else 14,color=NAVY,bold=True)
def add_chart(slide,chart_type,categories,series,left=.72,top=1.75,width=8.3,height=4.95,colors=None,number_format="0"):
    data=ChartData(); data.categories=list(categories)
    for name,values in series: data.add_series(name,list(values))
    chart=slide.shapes.add_chart(chart_type,Inches(left),Inches(top),Inches(width),Inches(height),data).chart; chart.chart_style=10; chart.has_legend=len(series)>1
    if chart.has_legend: chart.legend.position=XL_LEGEND_POSITION.BOTTOM; chart.legend.include_in_layout=False; chart.legend.font.size=Pt(8)
    chart.value_axis.has_major_gridlines=True; chart.value_axis.tick_labels.font.size=Pt(7.5); chart.value_axis.tick_labels.number_format=number_format; chart.category_axis.tick_labels.font.size=Pt(7); chart.category_axis.tick_label_skip=max(1,len(categories)//10)
    colors=colors or [BLUE,CYAN,AMBER]
    for index,item in enumerate(chart.series):
        item.format.line.color.rgb=colors[index%len(colors)]; item.format.line.width=Pt(1.7); item.format.fill.solid(); item.format.fill.fore_color.rgb=colors[index%len(colors)]
    if chart_type==XL_CHART_TYPE.COLUMN_CLUSTERED: chart.plots[0].gap_width=35
    return chart
def verify(input_dir,clean_path):
    for name,digest in HASHES.items():
        path=input_dir/name
        if not path.is_file() or sha(path)!=digest: raise ValueError(f"冻结输入错误：{name}")
    if not clean_path.is_file() or sha(clean_path)!=CLEAN_HASH: raise ValueError("上游清洗真值错误")
    raw=pd.read_csv(input_dir/"stock_data_raw.csv"); clean=pd.read_csv(clean_path)
    candidate=raw.dropna().sort_values("Date").reset_index(drop=True); candidate["Unnamed: 0"]=candidate["Unnamed: 0"].astype("int64"); candidate["Volume"]=candidate["Volume"].astype("int64")
    if len(raw)!=3740 or len(clean)!=3735 or not candidate.equals(clean): raise ValueError("原始数据与上游清洗真值不一致")
    return clean
def heat_color(value):
    value=max(-1,min(1,float(value))); base=np.array([255,255,255]); target=np.array([47,111,212]) if value>=0 else np.array([199,67,59]); mixed=(base*(1-abs(value)) + target*abs(value)).astype(int); return RGBColor(*mixed.tolist())
def build(input_dir,clean_path,output):
    df=verify(input_dir,clean_path); df["Date"]=pd.to_datetime(df["Date"],utc=True); df=df.sort_values("Date").reset_index(drop=True); df["Daily_Return"]=df["Close"].pct_change(); df["MA50"]=df["Close"].rolling(50).mean(); df["MA200"]=df["Close"].rolling(200).mean(); monthly=df.set_index("Date").resample("ME").last(); categories=[index.strftime("%Y-%m") for index in monthly.index]
    total_return=df["Close"].iloc[-1]/df["Close"].iloc[0]-1; annual_vol=df["Daily_Return"].std()*np.sqrt(252); ma50=df["MA50"].iloc[-1]; ma200=df["MA200"].iloc[-1]
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]

    slide=prs.slides.add_slide(blank); header(slide,"BOARD STOCK REVIEW  /  TREND","Closing Prices Over Time",1); add_chart(slide,XL_CHART_TYPE.LINE,categories,[("Close",monthly["Close"].tolist())],colors=[BLUE],number_format="$0"); note_card(slide,"Period","2010-01-04\n→ 2024-11-04",1.84,BLUE); note_card(slide,"Close","$30.95  →  $408.46",2.92,GREEN); note_card(slide,"Cumulative change",f"+{total_return*100:,.2f}%",4.0,GREEN); rect(slide,9.42,5.08,3.18,1.1,PALE,PALE); text(slide,"Board read",9.68,5.27,2.5,.17,size=8,color=BLUE,bold=True); text(slide,"Strong long-run appreciation; historical trend is not a forecast.",9.68,5.58,2.55,.42,size=10,color=NAVY,bold=True)

    ma50_monthly=[None if pd.isna(value) else float(value) for value in monthly["MA50"]]; ma200_monthly=[None if pd.isna(value) else float(value) for value in monthly["MA200"]]
    slide=prs.slides.add_slide(blank); header(slide,"TREND SIGNALS","Close Price with 50 and 200-Day Moving Averages",2); add_chart(slide,XL_CHART_TYPE.LINE,categories,[("Close",monthly["Close"].tolist()),("50-Day MA",ma50_monthly),("200-Day MA",ma200_monthly)],colors=[BLUE,CYAN,AMBER],number_format="$0"); note_card(slide,"Latest close","$408.46",1.84,BLUE); note_card(slide,"50-day MA",f"${ma50:,.2f}",2.92,CYAN); note_card(slide,"200-day MA",f"${ma200:,.2f}",4.0,AMBER); rect(slide,9.42,5.08,3.18,1.1,PALE,PALE); text(slide,"Board read",9.68,5.27,2.5,.17,size=8,color=BLUE,bold=True); text(slide,"Latest close is below both averages; monitor, not a buy/sell signal.",9.68,5.58,2.55,.42,size=10,color=NAVY,bold=True)

    slide=prs.slides.add_slide(blank); header(slide,"BOARD STOCK REVIEW  /  RELATIONSHIPS","Correlation Between Columns",3); variables=["Adj_Close","Close","High","Low","Open","Volume","Daily_Return"]; corr=df[variables].corr(); table=slide.shapes.add_table(8,8,Inches(.72),Inches(1.72),Inches(8.3),Inches(4.85)).table
    table.columns[0].width=Inches(1.35)
    for col in range(1,8): table.columns[col].width=Inches(.99)
    for col,value in enumerate(["",*variables]):
        cell=table.cell(0,col); cell.text=value.replace("Daily_Return","Daily\nReturn").replace("Adj_Close","Adj\nClose"); cell.fill.solid(); cell.fill.fore_color.rgb=NAVY
    for row,name in enumerate(variables,1):
        cell=table.cell(row,0); cell.text=name.replace("Daily_Return","Daily Return").replace("Adj_Close","Adj Close"); cell.fill.solid(); cell.fill.fore_color.rgb=NAVY
        for col,name2 in enumerate(variables,1):
            value=float(corr.loc[name,name2]); cell=table.cell(row,col); cell.text=f"{value:.2f}"; cell.fill.solid(); cell.fill.fore_color.rgb=heat_color(value)
    for row in range(8):
        for col in range(8):
            cell=table.cell(row,col); cell.margin_left=cell.margin_right=cell.margin_top=cell.margin_bottom=Inches(.02); cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment=PP_ALIGN.CENTER
                for run in paragraph.runs: run.font.name="Aptos"; run.font.size=Pt(7.2 if row==0 or col==0 else 8.5); run.font.bold=True; run.font.color.rgb=WHITE if row==0 or col==0 or abs(float(corr.iloc[row-1,col-1]))>.65 else INK
    note_card(slide,"Close vs volume",f"{corr.loc['Close','Volume']:.2f}",1.84,RED); note_card(slide,"Return vs volume",f"{corr.loc['Daily_Return','Volume']:.2f}",2.92,AMBER); note_card(slide,"Price fields","~1.00",4.0,BLUE); rect(slide,9.42,5.08,3.18,1.1,PALE,PALE); text(slide,"Board read",9.68,5.27,2.5,.17,size=8,color=BLUE,bold=True); text(slide,"Volume is negatively associated with price level here; correlation is not causation.",9.68,5.54,2.55,.5,size=9.5,color=NAVY,bold=True)

    returns=df["Daily_Return"].dropna().to_numpy(); return_counts,return_edges=np.histogram(returns,bins=30); return_labels=[f"{((return_edges[i]+return_edges[i+1])/2)*100:.1f}%" for i in range(len(return_counts))]
    slide=prs.slides.add_slide(blank); header(slide,"BOARD STOCK REVIEW  /  RISK","Daily Return Distribution",4); add_chart(slide,XL_CHART_TYPE.COLUMN_CLUSTERED,return_labels,[("Daily Return",return_counts.tolist())],colors=[CYAN],number_format="0"); note_card(slide,"Mean daily return",f"{df['Daily_Return'].mean()*100:.2f}%",1.84,GREEN); note_card(slide,"Annualized volatility",f"{annual_vol*100:.2f}%",2.92,AMBER); note_card(slide,"Observed range",f"{returns.min()*100:.2f}%  →  +{returns.max()*100:.2f}%",4.0,RED); rect(slide,9.42,5.08,3.18,1.1,PALE,PALE); text(slide,"Board read",9.68,5.27,2.5,.17,size=8,color=BLUE,bold=True); text(slide,"Most days cluster near zero, with material tail events defining risk.",9.68,5.58,2.55,.42,size=10,color=NAVY,bold=True)

    volumes=df["Volume"].to_numpy()/1e6; volume_counts,volume_edges=np.histogram(volumes,bins=30); volume_labels=[f"{(volume_edges[i]+volume_edges[i+1])/2:.0f}M" for i in range(len(volume_counts))]
    slide=prs.slides.add_slide(blank); header(slide,"BOARD STOCK REVIEW  /  MARKET INTEREST","Volume Distribution",5); add_chart(slide,XL_CHART_TYPE.COLUMN_CLUSTERED,volume_labels,[("Volume",volume_counts.tolist())],colors=[AMBER],number_format="0"); note_card(slide,"Median volume",f"{df['Volume'].median()/1e6:.2f}M",1.84,BLUE); note_card(slide,"Mean volume",f"{df['Volume'].mean()/1e6:.2f}M",2.92,GREEN); note_card(slide,"Maximum volume",f"{df['Volume'].max()/1e6:.2f}M",4.0,RED); rect(slide,9.42,5.08,3.18,1.1,PALE,PALE); text(slide,"Board read",9.68,5.27,2.5,.17,size=8,color=BLUE,bold=True); text(slide,"Right-skewed activity: a small number of sessions carry exceptional interest.",9.68,5.58,2.55,.42,size=10,color=NAVY,bold=True)
    output.parent.mkdir(parents=True,exist_ok=True); prs.save(output)
    if len(Presentation(output).slides)!=5: raise ValueError("页数不匹配")
    print(f"{output}\nsha256={sha(output)}")
def main():
    parser=argparse.ArgumentParser(); parser.add_argument("--input-dir",type=Path,required=True); parser.add_argument("--clean",type=Path,required=True); parser.add_argument("--output",type=Path,required=True); args=parser.parse_args(); build(args.input_dir,args.clean,args.output)
if __name__=="__main__": main()
