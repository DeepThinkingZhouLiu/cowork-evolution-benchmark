#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import shutil
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.util import Inches

CSV_NAME="stock_cleaned.csv"; PPTX_NAME="stock_analysis.pptx"; SCRIPT_NAME="stock_analysis_script.md"
def directory(root,name): path=root/name; path.mkdir(parents=True,exist_ok=True); return path
def copy_all(golden,target):
    target=directory(target.parent,target.name)
    for name in (CSV_NAME,PPTX_NAME,SCRIPT_NAME): shutil.copy2(golden/name,target/name)
    return target
def replace_on_slide(slide,old,new):
    count=0
    for shape in slide.shapes:
        if not getattr(shape,"has_text_frame",False): continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if old in run.text: run.text=run.text.replace(old,new); count+=1
    if not count: raise ValueError(old)
def replace_all(prs,old,new):
    count=0
    for slide in prs.slides:
        for shape in slide.shapes:
            if not getattr(shape,"has_text_frame",False): continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if old in run.text: run.text=run.text.replace(old,new); count+=1
    if not count: raise ValueError(old)
def mutate_ppt(golden,target,operation):
    target=copy_all(golden,target); prs=Presentation(golden/PPTX_NAME); operation(prs); prs.save(target/PPTX_NAME)
def mutate_script(golden,target,old,new):
    target=copy_all(golden,target); content=(golden/SCRIPT_NAME).read_text();
    if old not in content: raise ValueError(old)
    (target/SCRIPT_NAME).write_text(content.replace(old,new))
def chart_on(slide): return next(shape.chart for shape in slide.shapes if getattr(shape,"has_chart",False))
def change_chart_value(prs,slide_index,series_name,value_index,delta):
    chart=chart_on(prs.slides[slide_index]); data=ChartData(); data.categories=[str(value) for value in chart.plots[0].categories]
    found=False
    for series in chart.series:
        values=list(series.values)
        if series.name==series_name:
            values[value_index]=float(values[value_index])+delta; found=True
        data.add_series(series.name,values)
    if not found: raise ValueError(series_name)
    chart.replace_data(data)
def delete_slide(prs,index):
    ids=prs.slides._sldIdLst; slide_id=ids[index]; prs.part.drop_rel(slide_id.rId); ids.remove(slide_id)
def main():
    parser=argparse.ArgumentParser(); parser.add_argument("--golden-dir",type=Path,required=True); parser.add_argument("--outdir",type=Path,required=True); args=parser.parse_args(); golden=args.golden_dir
    directory(args.outdir,"empty"); output=directory(args.outdir,"corrupt"); (output/PPTX_NAME).write_bytes(b"bad\n"); shutil.copy2(golden/CSV_NAME,output/CSV_NAME); shutil.copy2(golden/SCRIPT_NAME,output/SCRIPT_NAME)
    output=directory(args.outdir,"ppt-only"); shutil.copy2(golden/PPTX_NAME,output/PPTX_NAME); output=directory(args.outdir,"csv-only"); shutil.copy2(golden/CSV_NAME,output/CSV_NAME); output=directory(args.outdir,"script-only"); shutil.copy2(golden/SCRIPT_NAME,output/SCRIPT_NAME)
    output=copy_all(golden,args.outdir/"valid-repackage")
    with zipfile.ZipFile(golden/PPTX_NAME) as source,zipfile.ZipFile(output/PPTX_NAME,"w",compression=zipfile.ZIP_STORED) as target:
        for info in source.infolist(): target.writestr(info.filename,source.read(info.filename))
    def alt_style(prs): prs.slides[0].background.fill.solid(); prs.slides[0].background.fill.fore_color.rgb=RGBColor(0xEE,0xF4,0xF8)
    mutate_ppt(golden,args.outdir/"valid-alt-style",alt_style)
    output=copy_all(golden,args.outdir/"wrong-cleaned-value"); rows=list(csv.reader((golden/CSV_NAME).open())); rows[1][3]=str(float(rows[1][3])+1);
    with (output/CSV_NAME).open("w",newline="") as file: csv.writer(file).writerows(rows)
    output=copy_all(golden,args.outdir/"unsorted-cleaned-data"); rows=list(csv.reader((golden/CSV_NAME).open())); rows[1],rows[2]=rows[2],rows[1]
    with (output/CSV_NAME).open("w",newline="") as file: csv.writer(file).writerows(rows)
    mutate_ppt(golden,args.outdir/"missing-slide",lambda prs:delete_slide(prs,4))
    mutate_ppt(golden,args.outdir/"wrong-slide-title",lambda prs:replace_on_slide(prs.slides[0],"Closing Prices Over Time","Stock Overview"))
    mutate_ppt(golden,args.outdir/"wrong-close-series",lambda prs:change_chart_value(prs,0,"Close",0,10))
    mutate_ppt(golden,args.outdir/"wrong-ma-series",lambda prs:change_chart_value(prs,1,"50-Day MA",178,10))
    def wrong_correlation(prs):
        table=next(shape.table for shape in prs.slides[2].shapes if getattr(shape,"has_table",False)); table.cell(2,6).text="0.37"
    mutate_ppt(golden,args.outdir/"wrong-correlation",wrong_correlation)
    mutate_ppt(golden,args.outdir/"wrong-return-histogram",lambda prs:change_chart_value(prs,3,"Daily Return",0,1))
    mutate_ppt(golden,args.outdir/"wrong-volume-histogram",lambda prs:change_chart_value(prs,4,"Volume",0,1))
    mutate_script(golden,args.outdir/"missing-script-greeting","Good morning, board members. Thank you for being present.","Presentation starts now.")
    mutate_script(golden,args.outdir/"wrong-script-volatility","25.66%","5.66%")
    mutate_script(golden,args.outdir/"missing-script-closing","Thank you for your time and for being present at the meeting.","End of presentation.")
    mutate_ppt(golden,args.outdir/"missing-sources",lambda prs:replace_all(prs,"Source:","Source omitted:"))
    image=args.outdir/"flat.png"; Image.new("RGB",(1600,900),"white").save(image)
    def flattened(prs):
        slide=prs.slides[0]
        for shape in list(slide.shapes): shape._element.getparent().remove(shape._element)
        slide.shapes.add_picture(str(image),0,0,width=prs.slide_width,height=prs.slide_height)
    mutate_ppt(golden,args.outdir/"flattened-first-slide",flattened)
    def marker(prs):
        box=prs.slides[-1].shapes.add_textbox(Inches(.1),Inches(.1),Inches(2),Inches(.2)); box.text="solution/Golden"
    mutate_ppt(golden,args.outdir/"hardcoded-marker",marker)
    output=copy_all(golden,args.outdir/"extra-output"); (output/"debug.txt").write_text("debug\n")
if __name__=="__main__": main()
