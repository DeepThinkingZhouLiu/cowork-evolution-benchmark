#!/usr/bin/env python3
"""TheAgentCompany 股价董事会演示任务的确定性 Judge。"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import posixpath
import re
import subprocess
import tempfile
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd
from lxml import etree
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE

CSV_NAME="stock_cleaned.csv"; PPTX_NAME="stock_analysis.pptx"; SCRIPT_NAME="stock_analysis_script.md"
TASK_ID="cowork-evo/theagentcompany-stock-analysis-board-deck"; RUBRIC_VERSION="theagentcompany-stock-analysis-ppt-adaptation-v0.1"; JUDGE_VERSION="theagentcompany-stock-analysis-deterministic-v0.1"
HASHES={"Original_Request.md":"621281d20fa3e0b4dbf6903ff871e035c99bb194a80b4fe3e1188f7f7954281b","stock_data_raw.csv":"5e83e1084cdebcfe78892cc345d38a8d584d2533e58efab722294a94edf1874c"}
TITLES=["Closing Prices Over Time","Close Price with 50 and 200-Day Moving Averages","Correlation Between Columns","Daily Return Distribution","Volume Distribution"]
VARIABLES=["Adj_Close","Close","High","Low","Open","Volume","Daily_Return"]
SUSPICIOUS=["solution/golden","judge-result","artifact_dir","file:///root","/root/solution","upstream-evaluator.py"]
SPECS={"R001":(6,"hurdle"),"R002":(16,"hurdle"),"R003":(14,"hurdle"),"R004":(16,"hurdle"),"R005":(14,"hurdle"),"R006":(14,"hurdle"),"R007":(10,"hurdle"),"R008":(8,"hurdle"),"R009":(2,"positive"),"R010":(-20,"penalty")}

def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def norm(value): return re.sub(r"\s+"," ",str(value).replace("\u00a0"," ").replace("\v"," ")).strip()
def locate(name):
    for path in (Path("/app/input_files")/name,Path(__file__).resolve().parents[1]/"data/input_files"/name):
        if path.is_file(): return path
    return None
def source_ok():
    missing=[]
    for name,digest in HASHES.items():
        path=locate(name)
        if path is None or sha(path)!=digest: missing.append(name)
    return not missing,missing
def canonical_clean():
    raw=pd.read_csv(locate("stock_data_raw.csv")); clean=raw.dropna().sort_values("Date").reset_index(drop=True); clean["Unnamed: 0"]=clean["Unnamed: 0"].astype("int64"); clean["Volume"]=clean["Volume"].astype("int64"); return raw,clean
def package(path):
    evidence={"broken":None,"missing":[],"macros":[],"external":[]}
    try:
        with zipfile.ZipFile(path) as archive:
            evidence["broken"]=archive.testzip(); names=set(archive.namelist()); evidence["macros"]=[name for name in names if name.lower().endswith("vbaproject.bin")]
            for relation in sorted(name for name in names if name.endswith(".rels")):
                root=etree.fromstring(archive.read(relation)); owner="" if relation=="_rels/.rels" else posixpath.join(relation.split("/_rels/",1)[0],relation.split("/_rels/",1)[1][:-5]); base=posixpath.dirname(owner)
                for rel in root:
                    target=rel.attrib.get("Target","")
                    if rel.attrib.get("TargetMode")=="External": evidence["external"].append(target); continue
                    resolved=target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base,target))
                    if resolved not in names: evidence["missing"].append(f"{relation}->{resolved}")
        return evidence["broken"] is None and not evidence["missing"] and not evidence["macros"],evidence
    except Exception as exc: evidence["exception"]=str(exc); return False,evidence
def slide_text(slide): return norm("\n".join(shape.text for shape in slide.shapes if getattr(shape,"has_text_frame",False)))
def charts(slide): return [shape.chart for shape in slide.shapes if getattr(shape,"has_chart",False)]
def tables(slide): return [shape.table for shape in slide.shapes if getattr(shape,"has_table",False)]
def categories(chart): return [str(value) for value in chart.plots[0].categories]
def series_map(chart): return {series.name:list(series.values) for series in chart.series}
def close_values(actual,expected,atol=1e-7):
    if len(actual)!=len(expected): return False
    for left,right in zip(actual,expected):
        right_missing=right is None or (isinstance(right,float) and math.isnan(right)); left_missing=left is None or (isinstance(left,float) and math.isnan(left))
        if right_missing or left_missing:
            if right_missing!=left_missing: return False
        elif not math.isclose(float(left),float(right),rel_tol=1e-9,abs_tol=atol): return False
    return True
def render(path,out):
    out.mkdir(parents=True,exist_ok=True); profile=out/"profile"; profile.mkdir(exist_ok=True); process=subprocess.run(["soffice",f"-env:UserInstallation=file://{profile}","--headless","--convert-to","pdf","--outdir",str(out),str(path)],stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,timeout=300); pdf=out/f"{path.stem}.pdf"
    if process.returncode or not pdf.is_file(): return 0,process.stdout[-500:]
    info=subprocess.run(["pdfinfo",str(pdf)],stdout=subprocess.PIPE,text=True).stdout; match=re.search(r"^Pages:\s+(\d+)",info,re.M); return int(match.group(1)) if match else 0,process.stdout[-500:]
def item(cid,passed,evidence):
    weight,kind=SPECS[cid]; return {"criterion_id":cid,"score":1.0 if passed else 0.0,"evidence":evidence,"method":"deterministic","raw":{"weight":weight,"type":kind,"passed":bool(passed)}}
def assemble(paths,criteria):
    hurdles=[entry["criterion_id"] for entry in criteria if entry["raw"]["type"]=="hurdle" and not entry["raw"]["passed"]]; total=sum(entry["raw"]["weight"] for entry in criteria if entry["raw"]["weight"]>0); earned=sum(entry["raw"]["weight"] for entry in criteria if entry["raw"]["weight"]>0 and entry["raw"]["passed"]); penalty=sum(abs(entry["raw"]["weight"]) for entry in criteria if entry["raw"]["weight"]<0 and not entry["raw"]["passed"]); reward=0 if hurdles else max(0,min(1,(earned-penalty)/total)); return {"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":[sha(path) for path in paths if path.is_file()],"criterion_results":criteria,"reward":reward,"passed":reward>=.8 and not hurdles,"hurdle_failures":hurdles}
def evaluate(csv_path,pptx_path,script_path,artifact):
    src,source_evidence=source_ok()
    try: prs=Presentation(pptx_path)
    except Exception as exc: return assemble([csv_path,pptx_path,script_path],[item(cid,False,f"打开失败:{exc}") for cid in SPECS])
    pkg,pkg_evidence=package(pptx_path); ratio=prs.slide_width/prs.slide_height if prs.slide_height else 0; r1=src and pkg and csv_path.is_file() and script_path.is_file() and len(prs.slides)==5 and abs(ratio-16/9)<.02; criteria=[item("R001",r1,f"pages={len(prs.slides)} ratio={ratio} source={source_evidence} package={json.dumps(pkg_evidence,ensure_ascii=False)}")]
    raw,expected=canonical_clean(); clean_errors=[]
    try:
        actual=pd.read_csv(csv_path)
        if not actual.equals(expected):
            clean_errors.append(f"equals=False shape={actual.shape} dtypes={actual.dtypes.astype(str).to_dict()}")
        if actual.isna().any().any(): clean_errors.append("nan-present")
        if list(actual.columns)!=list(expected.columns): clean_errors.append("columns")
        if not actual["Date"].is_monotonic_increasing: clean_errors.append("date-order")
    except Exception as exc:
        actual=None; clean_errors.append(f"parse:{exc}")
    r2=len(raw)==3740 and len(expected)==3735 and not clean_errors; criteria.append(item("R002",r2,f"raw_rows={len(raw)} expected_rows={len(expected)} errors={clean_errors}"))
    df=expected.copy(); df["Date"]=pd.to_datetime(df["Date"],utc=True); df["Daily_Return"]=df["Close"].pct_change(); df["MA50"]=df["Close"].rolling(50).mean(); df["MA200"]=df["Close"].rolling(200).mean(); monthly=df.set_index("Date").resample("ME").last(); expected_categories=[index.strftime("%Y-%m") for index in monthly.index]; texts=[slide_text(slide) for slide in prs.slides]
    chart1=charts(prs.slides[0]) if len(prs.slides)>0 else []; map1=series_map(chart1[0]) if len(chart1)==1 else {}; r3=len(chart1)==1 and chart1[0].chart_type==XL_CHART_TYPE.LINE and categories(chart1[0])==expected_categories and set(map1)=={"Close"} and close_values(map1.get("Close",[]),monthly["Close"].tolist()) and all(value.lower() in texts[0].lower() for value in [TITLES[0],"2010-01-04","2024-11-04","$30.95","$408.46","+1,219.74%","not a forecast"]); criteria.append(item("R003",r3,f"charts={len(chart1)} type={chart1[0].chart_type if chart1 else None} series={list(map1)} categories={len(categories(chart1[0])) if chart1 else 0}"))
    chart2=charts(prs.slides[1]) if len(prs.slides)>1 else []; map2=series_map(chart2[0]) if len(chart2)==1 else {}; ma50=[None if pd.isna(value) else float(value) for value in monthly["MA50"]]; ma200=[None if pd.isna(value) else float(value) for value in monthly["MA200"]]; r4=len(chart2)==1 and chart2[0].chart_type==XL_CHART_TYPE.LINE and categories(chart2[0])==expected_categories and set(map2)=={"Close","50-Day MA","200-Day MA"} and close_values(map2.get("Close",[]),monthly["Close"].tolist()) and close_values(map2.get("50-Day MA",[]),ma50) and close_values(map2.get("200-Day MA",[]),ma200) and all(value.lower() in texts[1].lower() for value in [TITLES[1],"$408.46","$420.52","$420.88","below both averages","not a buy/sell signal"]); criteria.append(item("R004",r4,f"charts={len(chart2)} type={chart2[0].chart_type if chart2 else None} series={list(map2)}"))
    table_list=tables(prs.slides[2]) if len(prs.slides)>2 else []; table_errors=[]; corr=df[VARIABLES].corr()
    if len(table_list)!=1 or len(table_list[0].rows)!=8 or len(table_list[0].columns)!=8: table_errors.append("shape")
    else:
        table=table_list[0]; expected_labels=[name.replace("_"," ").lower() for name in VARIABLES]; row_labels=[norm(table.cell(row,0).text).lower() for row in range(1,8)]; col_labels=[norm(table.cell(0,col).text).lower() for col in range(1,8)]
        if row_labels!=expected_labels or col_labels!=expected_labels: table_errors.append(f"labels:{row_labels}:{col_labels}")
        for row,name in enumerate(VARIABLES,1):
            for col,name2 in enumerate(VARIABLES,1):
                try: value=float(norm(table.cell(row,col).text))
                except Exception: table_errors.append(f"parse:{row}:{col}"); continue
                if not math.isclose(value,round(float(corr.loc[name,name2]),2),abs_tol=.005): table_errors.append(f"value:{name}:{name2}:{value}")
    r5=not table_errors and all(value.lower() in texts[2].lower() for value in [TITLES[2],"close vs volume","-0.37","return vs volume","-0.07","correlation is not causation"]); criteria.append(item("R005",r5,f"tables={len(table_list)} errors={table_errors[:10]}"))
    returns=df["Daily_Return"].dropna().to_numpy(); return_counts,return_edges=np.histogram(returns,bins=30); return_labels=[f"{((return_edges[index]+return_edges[index+1])/2)*100:.1f}%" for index in range(30)]; volumes=df["Volume"].to_numpy()/1e6; volume_counts,volume_edges=np.histogram(volumes,bins=30); volume_labels=[f"{(volume_edges[index]+volume_edges[index+1])/2:.0f}M" for index in range(30)]; histogram_errors=[]
    for page,name,labels,counts,title,facts in [(3,"Daily Return",return_labels,return_counts.tolist(),TITLES[3],["0.08%","25.66%","-14.74%","+14.22%","tail events"]),(4,"Volume",volume_labels,volume_counts.tolist(),TITLES[4],["30.38M","36.32M","319.32M","Right-skewed"])]:
        page_charts=charts(prs.slides[page]) if len(prs.slides)>page else []
        if len(page_charts)!=1 or page_charts[0].chart_type!=XL_CHART_TYPE.COLUMN_CLUSTERED: histogram_errors.append(f"page={page+1}:chart"); continue
        mapping=series_map(page_charts[0])
        if categories(page_charts[0])!=labels or set(mapping)!={name} or not close_values(mapping[name],counts,0): histogram_errors.append(f"page={page+1}:data")
        if not all(value.lower() in texts[page].lower() for value in [title,*facts]): histogram_errors.append(f"page={page+1}:text")
    r6=not histogram_errors; criteria.append(item("R006",r6,f"errors={histogram_errors}"))
    try: script=norm(script_path.read_text())
    except Exception as exc: script=f"parse-error:{exc}"
    script_required=["Good morning, board members","Thank you for being present","3,740 raw rows","five invalid","3,735 dated observations","2010","2024","Slide 1","Slide 2","Slide 3","Slide 4","Slide 5",*TITLES,"1,219.74%","$420.52","$420.88","-0.37","-0.07","25.66%","-14.74%","+14.22%","30.38 million","319.32 million","not investment advice","correlation does not establish causation","Thank you for your time"]
    script_missing=[value for value in script_required if value.lower() not in script.lower()]; r7=not script_missing; criteria.append(item("R007",r7,f"missing={script_missing}"))
    alltext="\n".join(texts); title_errors=[f"page={index+1}:{title}" for index,title in enumerate(TITLES) if index>=len(texts) or title.lower() not in texts[index].lower()]; citations=sum(text.count("Source:") for text in texts); source_tags=sum("stock_cleaned.csv" in text and "stock_data_raw.csv" in text for text in texts); markers=[value for value in SUSPICIOUS if value in alltext.lower()]; picture_ids=[shape.shape_id for slide in prs.slides for shape in slide.shapes if shape.shape_type==MSO_SHAPE_TYPE.PICTURE]; native_charts=sum(len(charts(slide)) for slide in prs.slides); native_tables=sum(len(tables(slide)) for slide in prs.slides); native_text=sum(getattr(shape,"has_text_frame",False) for slide in prs.slides for shape in slide.shapes)
    with zipfile.ZipFile(pptx_path) as archive: media=[name for name in archive.namelist() if name.startswith("ppt/media/") and not name.endswith("/")]
    r8=not title_errors and citations==5 and source_tags==5 and not markers and not picture_ids and not media and native_charts==4 and native_tables==1 and native_text>=40 and not pkg_evidence.get("macros") and not pkg_evidence.get("external"); criteria.append(item("R008",r8,f"title_errors={title_errors} citations={citations} source_tags={source_tags} markers={markers} pictures={picture_ids} media={media} charts={native_charts} tables={native_tables} text={native_text}"))
    pages=0; log=""
    try:
        with tempfile.TemporaryDirectory(prefix="tac-stock-render-") as directory: pages,log=render(pptx_path,Path(directory))
    except Exception as exc: log=str(exc)
    delivered=sorted(str(path.relative_to(artifact)) for path in artifact.rglob("*") if path.is_file()); expected_files=sorted([CSV_NAME,PPTX_NAME,SCRIPT_NAME]); r9=pages==5 and delivered==expected_files and pkg; criteria.append(item("R009",r9,f"pages={pages} files={delivered} log={log}"))
    severe=not all([r2,r3,r4,r5,r6,r7,r8]); criteria.append(item("R010",not severe,f"severe={severe}")); return assemble([csv_path,pptx_path,script_path],criteria)
def main():
    parser=argparse.ArgumentParser(); parser.add_argument("--artifact-dir",type=Path,required=True); parser.add_argument("--result",type=Path,required=True); parser.add_argument("--reward-file",type=Path,required=True); args=parser.parse_args(); paths=[args.artifact_dir/CSV_NAME,args.artifact_dir/PPTX_NAME,args.artifact_dir/SCRIPT_NAME]
    result=evaluate(*paths,args.artifact_dir) if all(path.is_file() for path in paths) else assemble(paths,[item(cid,False,f"缺失交付文件：{CSV_NAME} / {PPTX_NAME} / {SCRIPT_NAME}") for cid in SPECS]); args.result.parent.mkdir(parents=True,exist_ok=True); args.reward_file.parent.mkdir(parents=True,exist_ok=True); args.result.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n"); args.reward_file.write_text(f"{result['reward']:.6f}\n"); print(json.dumps(result,ensure_ascii=False,indent=2))
if __name__=="__main__": main()
