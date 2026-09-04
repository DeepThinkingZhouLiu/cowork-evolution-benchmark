#!/usr/bin/env python3
"""TheAgentCompany finance-budget-variance 的逐题参考解法。"""

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

EXPECTED_HASHES = {
    "budget.xlsx": "c9578fd151590966ee1003559132fcf81ede7aa249393ec8d9e3a57ce063607d",
    "actual_spending.xlsx": "060d0ecadab3863cc71fcb838e10e84cca73c81beb4d407b063ee410f5e2d139",
}
HEADERS = ["Flagged_Variances", "Department", "Category", "Month", "Budgeted_Amount", "Actual_Amount", "Variance_Amount", "Variance_Percentage"]


def digest(path: Path) -> str:
    value = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            value.update(block)
    return value.hexdigest()


def rows(path: Path) -> list[tuple]:
    workbook = load_workbook(path, data_only=True)
    values = list(workbook.active.iter_rows(min_row=2, values_only=True))
    workbook.close()
    return values


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    for name, expected in EXPECTED_HASHES.items():
        if digest(args.input_dir / name) != expected:
            raise ValueError(f"输入哈希不匹配：{name}")
    budget_rows = rows(args.input_dir / "budget.xlsx")
    actual_rows = rows(args.input_dir / "actual_spending.xlsx")
    actual = {(department, category, month): amount for department, category, month, amount in actual_rows}
    if len(budget_rows) != 420 or set(actual) != {(row[0], row[1], row[2]) for row in budget_rows}:
        raise ValueError("预算和实际联结键不完整")
    flagged = []
    for department, category, month, budget in budget_rows:
        actual_amount = actual[(department, category, month)]
        variance = actual_amount - budget
        percentage = variance / budget * 100
        if abs(percentage) > 10 and abs(variance) > 5000:
            flagged.append([f"{department}_{category}_{month:%Y%m}", department, category, month, budget, actual_amount, variance, percentage])

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Flagged_Variances"
    sheet.append(HEADERS)
    for row in flagged:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        row[3].number_format = "yyyy-mm-dd"
        for cell in row[4:7]:
            cell.number_format = "$#,##0.00;[Red]-$#,##0.00"
        row[7].number_format = "0.00%;[Red]-0.00%"
        row[7].value = row[7].value / 100
    widths = [36, 16, 16, 14, 18, 18, 18, 20]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[__import__("openpyxl").utils.get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:H{sheet.max_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:1"
    sheet.sheet_view.showGridLines = False
    table = Table(displayName="FlaggedVariances", ref=f"A1:H{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output_dir / "flagged_variances.xlsx")
    workbook.close()


if __name__ == "__main__":
    main()
