#!/usr/bin/env python3
"""确定性 Judge：从两份输入独立联结并重算重大预算差异。"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

TASK_ID = "cowork-evo/agentcompany-budget-variance-analysis"
RUBRIC_VERSION = "agentcompany-budget-variance-analysis-v0.1"
JUDGE_VERSION = "agentcompany-budget-variance-analysis-deterministic-v0.1"
OUTPUT_NAME = "flagged_variances.xlsx"
HEADERS = ["Flagged_Variances", "Department", "Category", "Month", "Budgeted_Amount", "Actual_Amount", "Variance_Amount", "Variance_Percentage"]
EXPECTED_HASHES = {"budget.xlsx": "c9578fd151590966ee1003559132fcf81ede7aa249393ec8d9e3a57ce063607d", "actual_spending.xlsx": "060d0ecadab3863cc71fcb838e10e84cca73c81beb4d407b063ee410f5e2d139"}
WEIGHTS = {"R001": 0.10, "R002": 0.25, "R003": 0.15, "R004": 0.20, "R005": 0.10, "R006": 0.05, "R007": 0.10, "R008": 0.05}


def digest(path: Path) -> str:
    value = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            value.update(block)
    return value.hexdigest()


def criterion(cid: str, score: bool | float, evidence: str) -> dict:
    return {"criterion_id": cid, "score": float(score), "evidence": evidence, "method": "deterministic"}


def source_rows(path: Path) -> list[tuple]:
    workbook = load_workbook(path, data_only=True)
    values = list(workbook.active.iter_rows(min_row=2, values_only=True))
    workbook.close()
    return values


def yyyymm(value: object) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y%m")
    text = str(value).strip()
    digits = re.sub(r"[^0-9]", "", text)
    return digits[:6] if len(digits) >= 6 else text


def as_percent(value: object, number_format: str) -> float:
    raw = float(value)
    return raw * 100 if "%" in number_format else raw


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    input_dir = Path(__file__).parents[1] / "data/input_files"
    output_path = args.artifact_dir / OUTPUT_NAME
    results: list[dict] = []
    try:
        if not output_path.is_file() or not zipfile.is_zipfile(output_path):
            raise ValueError("输出缺失或不是有效 OOXML")
        budget_rows = source_rows(input_dir / "budget.xlsx")
        actual_rows = source_rows(input_dir / "actual_spending.xlsx")
        actual = {(department, category, month): amount for department, category, month, amount in actual_rows}
        expected = []
        for department, category, month, budget in budget_rows:
            actual_amount = actual[(department, category, month)]
            variance = actual_amount - budget
            percentage = variance / budget * 100
            if abs(percentage) > 10 and abs(variance) > 5000:
                expected.append({"id": f"{department}_{category}_{month:%Y%m}", "department": department, "category": category, "month": month.strftime("%Y%m"), "budget": float(budget), "actual": float(actual_amount), "variance": float(variance), "percentage": float(percentage)})
        workbook = load_workbook(output_path, data_only=False)
        if "Flagged_Variances" not in workbook.sheetnames:
            raise ValueError("缺少 Flagged_Variances 工作表")
        sheet = workbook["Flagged_Variances"]
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
        missing_headers = [name for name in HEADERS if name not in headers]
        results.append(criterion("R001", not missing_headers, f"缺少列={missing_headers}"))
        index = {name: headers.index(name) + 1 for name in HEADERS}
        output_rows = []
        for row_number in range(2, sheet.max_row + 1):
            identifier = sheet.cell(row_number, index["Flagged_Variances"]).value
            if identifier is None:
                continue
            pct_cell = sheet.cell(row_number, index["Variance_Percentage"])
            output_rows.append({"id": str(identifier).strip(), "department": str(sheet.cell(row_number, index["Department"]).value).strip(), "category": str(sheet.cell(row_number, index["Category"]).value).strip(), "month": yyyymm(sheet.cell(row_number, index["Month"]).value), "budget": float(sheet.cell(row_number, index["Budgeted_Amount"]).value), "actual": float(sheet.cell(row_number, index["Actual_Amount"]).value), "variance": float(sheet.cell(row_number, index["Variance_Amount"]).value), "percentage": as_percent(pct_cell.value, pct_cell.number_format)})
        expected_ids = [row["id"] for row in expected]
        output_ids = [row["id"] for row in output_rows]
        results.append(criterion("R002", output_ids == expected_ids, f"期望={expected_ids}; 输出={output_ids}"))
        traces_ok = len(output_rows) == len(expected) and all(out[key] == exp[key] if key in {"department", "category", "month"} else abs(out[key] - exp[key]) < 0.01 for out, exp in zip(output_rows, expected) for key in ("department", "category", "month", "budget", "actual"))
        results.append(criterion("R003", traces_ok, f"来源字段一致={traces_ok}"))
        calculations_ok = len(output_rows) == len(expected) and all(abs(out["variance"] - exp["variance"]) <= 0.01 and abs(out["percentage"] - exp["percentage"]) <= 0.01 for out, exp in zip(output_rows, expected))
        results.append(criterion("R004", calculations_ok, "差异金额和百分比按输入重算"))
        strict_ok = len(expected) == 8 and output_ids == expected_ids and all(abs(row["percentage"]) > 10 and abs(row["variance"]) > 5000 for row in output_rows)
        results.append(criterion("R005", strict_ok, f"重算重大差异数={len(expected)}"))
        format_ok = len(output_ids) == len(set(output_ids)) and all(re.fullmatch(r"[A-Za-z]+_[A-Za-z]+_[0-9]{6}", value) for value in output_ids)
        results.append(criterion("R006", format_ok and output_ids == expected_ids, "标识格式、唯一性和输入顺序检查"))
        table_ok = bool(sheet.tables) and sheet.freeze_panes == "A2" and sheet.auto_filter.ref is not None and all((sheet.column_dimensions[letter].width or 0) >= 12 for letter in ("A", "B", "C", "D"))
        page_ok = sheet.page_setup.orientation == "landscape" and sheet.page_setup.fitToWidth == 1
        results.append(criterion("R007", table_ok and page_ok and sheet.protection.sheet is not True, f"表格={bool(sheet.tables)}; 冻结={sheet.freeze_panes}; 筛选={sheet.auto_filter.ref}; 横向单页宽={page_ok}"))
        hashes_ok = all(digest(input_dir / name) == expected_hash for name, expected_hash in EXPECTED_HASHES.items())
        key_alignment = len(budget_rows) == len(actual_rows) == 420 and {(row[0], row[1], row[2]) for row in budget_rows} == set(actual)
        results.append(criterion("R008", hashes_ok and key_alignment, f"哈希={hashes_ok}; 联结键对齐={key_alignment}"))
        workbook.close()
    except Exception as exc:
        results = [criterion("R001", False, f"Judge 读取失败：{exc}")]
    by_id = {item["criterion_id"]: item["score"] for item in results}
    reward = sum(WEIGHTS[cid] * by_id.get(cid, 0.0) for cid in WEIGHTS)
    passed = reward >= 0.8 and all(by_id.get(cid) == 1.0 for cid in ("R001", "R002", "R003", "R004", "R005"))
    payload = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [digest(output_path)] if output_path.is_file() else [], "criterion_results": results, "reward": reward, "passed": passed}
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.write_text(f"{reward:.6f}\n", encoding="utf-8")
    raise SystemExit(0 if passed else 1)


if __name__ == "__main__":
    main()
