#!/usr/bin/env python3
"""从固定输入复制工作表并完成任务要求的格式化。"""

from __future__ import annotations

import argparse
import hashlib
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

INPUT_NAME = "respondent_counts_input.xlsx"
OUTPUT_NAME = "respondent_counts.xlsx"
SHEET_NAME = "unformatted"
FORMATTED_NAME = "formatted"
EXPECTED_SHA256 = "be71871d79156a2106b085cf29a20e08ee39bf305c82022003f9533d33ffc043"
TARGET_CELLS = {
    "A3", "B3", "B5", "C3", "C4", "D4", "D5", "E5", "F3", "F4",
    "G4", "G5", "H5", "I3", "I4", "J4", "J5", "K5",
}


def digest(path: Path) -> str:
    value = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            value.update(block)
    return value.hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()

    input_path = args.input_dir / INPUT_NAME
    if digest(input_path) != EXPECTED_SHA256:
        raise ValueError("输入工作簿 SHA256 不匹配")

    workbook = load_workbook(input_path)
    if workbook.sheetnames != [SHEET_NAME]:
        raise ValueError(f"输入工作表应仅有 {SHEET_NAME}，实际为 {workbook.sheetnames}")
    source = workbook[SHEET_NAME]
    formatted = workbook.copy_worksheet(source)
    formatted.title = FORMATTED_NAME

    # 只在副本上修改水平对齐和指定表头填充，内容及其它样式来自输入副本。
    for row in formatted.iter_rows(min_row=1, max_row=source.max_row, min_col=1, max_col=source.max_column):
        for cell in row:
            if cell.value is not None:
                alignment = copy(cell.alignment)
                alignment.horizontal = "center"
                cell.alignment = alignment

    skyblue = PatternFill(fill_type="solid", fgColor="FF87CEEB")
    for coordinate in TARGET_CELLS:
        formatted[coordinate].fill = copy(skyblue)

    args.output_dir.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output_dir / OUTPUT_NAME)
    workbook.close()


if __name__ == "__main__":
    main()
