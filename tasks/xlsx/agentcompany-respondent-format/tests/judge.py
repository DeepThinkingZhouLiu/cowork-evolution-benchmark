#!/usr/bin/env python3
"""从输入独立检查 respondent counts 格式化任务。"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import subprocess
import tempfile
import zipfile
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

INPUT_NAME = "respondent_counts_input.xlsx"
OUTPUT_NAME = "respondent_counts.xlsx"
SOURCE_SHEET = "unformatted"
FORMATTED_SHEET = "formatted"
EXPECTED_HASH = "be71871d79156a2106b085cf29a20e08ee39bf305c82022003f9533d33ffc043"
TARGET_CELLS = {
    "A3", "B3", "B5", "C3", "C4", "D4", "D5", "E5", "F3", "F4",
    "G4", "G5", "H5", "I3", "I4", "J4", "J5", "K5",
}


def digest(path: Path) -> str:
    value = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            value.update(block)
    return value.hexdigest()


def color_key(color) -> tuple:
    if color is None:
        return (None, None, None, None, None)
    return (color.type, color.rgb, color.indexed, color.theme, color.tint)


def side_key(side) -> tuple:
    return (side.style, color_key(side.color))


def fill_key(cell) -> tuple:
    return (cell.fill.fill_type, color_key(cell.fill.fgColor), color_key(cell.fill.bgColor))


def style_key(cell, *, include_fill: bool = True, include_alignment: bool = True) -> tuple:
    font = cell.font
    border = cell.border
    protection = cell.protection
    alignment = cell.alignment
    result = (
        (font.name, font.sz, font.bold, font.italic, font.underline, font.strike, color_key(font.color)),
        (side_key(border.left), side_key(border.right), side_key(border.top), side_key(border.bottom), side_key(border.diagonal)),
        cell.number_format,
        (protection.locked, protection.hidden),
    )
    if include_fill:
        result += ((cell.fill.fill_type, color_key(cell.fill.fgColor), color_key(cell.fill.bgColor)),)
    if include_alignment:
        result += ((alignment.horizontal, alignment.vertical, alignment.textRotation, alignment.wrap_text or False, alignment.shrink_to_fit or False, alignment.indent),)
    return result


def criterion(cid: str, passed: bool, evidence: str) -> dict:
    return {"criterion_id": cid, "score": 1.0 if passed else 0.0, "passed": bool(passed), "evidence": evidence, "method": "deterministic"}


def values(ws):
    return [[ws.cell(r, c).value for c in range(1, ws.max_column + 1)] for r in range(1, ws.max_row + 1)]


def dimensions(ws) -> tuple:
    widths = tuple((letter, ws.column_dimensions[letter].width, ws.column_dimensions[letter].hidden) for letter in "ABCDEFGHIJK")
    heights = tuple((row, ws.row_dimensions[row].height, ws.row_dimensions[row].hidden) for row in range(1, 22))
    return (widths, heights, ws.freeze_panes, ws.auto_filter.ref, tuple(sorted(str(x) for x in ws.merged_cells.ranges)), ws.sheet_state)


def render_ok(path: Path) -> tuple[bool, str]:
    libreoffice = shutil.which("libreoffice")
    if libreoffice is None:
        return False, "libreoffice 不可用"
    with tempfile.TemporaryDirectory(prefix="c082-render-") as temp:
        out = Path(temp)
        profile = out / "profile"
        profile.mkdir()
        result = subprocess.run(
            [libreoffice, "--headless", f"-env:UserInstallation={profile.resolve().as_uri()}", "--convert-to", "pdf", "--outdir", str(out), str(path)],
            capture_output=True, text=True, timeout=180, check=False,
        )
        pdf = out / (path.stem + ".pdf")
        return result.returncode == 0 and pdf.is_file() and pdf.stat().st_size > 1000, (result.stderr or result.stdout or "")[-500:]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    input_path = Path(__file__).parents[1] / "data/input_files" / INPUT_NAME
    output_path = args.artifact_dir / OUTPUT_NAME
    results = []
    source = candidate = None
    try:
        valid = output_path.is_file() and zipfile.is_zipfile(output_path)
        if not valid:
            raise ValueError("输出缺失或不是有效 OOXML")
        source_wb = load_workbook(input_path, data_only=False)
        candidate_wb = load_workbook(output_path, data_only=False)
        source = source_wb[SOURCE_SHEET]
        candidate = candidate_wb[SOURCE_SHEET] if SOURCE_SHEET in candidate_wb.sheetnames else None
        formatted = candidate_wb[FORMATTED_SHEET] if FORMATTED_SHEET in candidate_wb.sheetnames else None
        results.append(criterion("R001", candidate_wb.sheetnames == [SOURCE_SHEET, FORMATTED_SHEET], f"工作表={candidate_wb.sheetnames}"))

        unchanged = candidate is not None and values(source) == values(candidate) and dimensions(source) == dimensions(candidate)
        # 比较输入与输出 unformatted 的所有可见样式属性，允许 OOXML 写入时的 None/False 归一化。
        if unchanged:
            for row in range(1, 22):
                for col in range(1, 12):
                    if style_key(source.cell(row, col)) != style_key(candidate.cell(row, col)):
                        unchanged = False
                        break
                if not unchanged:
                    break
        results.append(criterion("R002", unchanged, f"unformatted 内容/结构/样式一致={unchanged}"))

        copied = formatted is not None and formatted.max_row == source.max_row and formatted.max_column == source.max_column and values(formatted) == values(source) and tuple(sorted(str(x) for x in formatted.merged_cells.ranges)) == tuple(sorted(str(x) for x in source.merged_cells.ranges))
        results.append(criterion("R003", copied, f"formatted 复制一致={copied}"))

        fill_ok = False
        if formatted is not None:
            fill_ok = all(formatted[ref].fill.fill_type == "solid" and formatted[ref].fill.fgColor.type == "rgb" and formatted[ref].fill.fgColor.rgb == "FF87CEEB" for ref in TARGET_CELLS)
        results.append(criterion("R004", fill_ok, f"目标填充正确={fill_ok}"))

        centered = False
        if formatted is not None:
            populated = [cell for row in formatted.iter_rows(min_row=1, max_row=source.max_row, min_col=1, max_col=source.max_column) for cell in row if cell.value is not None]
            centered = bool(populated) and all(cell.alignment.horizontal == "center" for cell in populated)
        results.append(criterion("R005", centered, f"非空单元格居中={centered}"))

        collateral = False
        if formatted is not None:
            collateral = True
            for row in range(1, 22):
                for col in range(1, 12):
                    before, after = source.cell(row, col), formatted.cell(row, col)
                    if before.value != after.value or style_key(before, include_fill=False, include_alignment=False) != style_key(after, include_fill=False, include_alignment=False):
                        collateral = False
                        break
                    if fill_key(before) != fill_key(after) and after.coordinate not in TARGET_CELLS:
                        collateral = False
                        break
                if not collateral:
                    break
        results.append(criterion("R006", collateral, f"非目标内容和样式无额外变化={collateral}"))

        native = formatted is not None and all(ws.protection.sheet is not True for ws in candidate_wb.worksheets)
        media = []
        try:
            with zipfile.ZipFile(output_path) as archive:
                media = [name for name in archive.namelist() if name.startswith("xl/media/")]
        except zipfile.BadZipFile:
            native = False
        results.append(criterion("R007", native and not media, f"可编辑={native}; 媒体={media}"))

        hashes = digest(input_path) == EXPECTED_HASH
        rendered, render_evidence = render_ok(output_path)
        results.append(criterion("R008", hashes and rendered, f"输入哈希={hashes}; 渲染={rendered}; {render_evidence}"))
        source_wb.close()
        candidate_wb.close()
    except Exception as exc:
        results = [criterion("R001", False, f"Judge 读取失败：{exc}")]

    weights = {"R001": 0.15, "R002": 0.20, "R003": 0.15, "R004": 0.20, "R005": 0.15, "R006": 0.08, "R007": 0.04, "R008": 0.03}
    by_id = {item["criterion_id"]: item["score"] for item in results}
    reward = sum(weights[cid] * by_id.get(cid, 0.0) for cid in weights)
    passed = reward >= 0.8 and all(by_id.get(cid) == 1.0 for cid in ("R001", "R002", "R003", "R004", "R005"))
    payload = {"task_id": "cowork-evo/agentcompany-respondent-format", "rubric_version": "agentcompany-respondent-format-v0.1", "judge_version": "agentcompany-respondent-format-deterministic-v0.1", "artifact_sha256": [digest(output_path)] if output_path.is_file() else [], "criterion_results": results, "reward": reward, "passed": passed}
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.write_text(f"{reward:.6f}\n", encoding="utf-8")
    raise SystemExit(0 if passed else 1)


if __name__ == "__main__":
    main()
