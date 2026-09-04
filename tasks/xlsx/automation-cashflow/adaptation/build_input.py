"""Create the solver-visible workbook from the fixed AutomationBench fixture."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).parents[1]

def style_sheet(ws, end_col, end_row):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{end_col}{end_row}"
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    tab = Table(displayName=f"{ws.title.replace(' ', '')}Table", ref=f"A1:{end_col}{end_row}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)

def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Bank Balance"
    ws.append(["Account", "Balance", "As Of"])
    ws.append(["Operating Account", 200000, "2026-02-10"])
    ws["B2"].number_format = '$#,##0;[Red]-$#,##0'
    style_sheet(ws, "C", 2)
    ws.column_dimensions["A"].width = 22; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 14

    ws = wb.create_sheet("AR Due")
    ws.append(["Customer", "Amount", "Collection Probability", "Status"])
    for row in [
        ["NovaTech Solutions", 40000, 100, "Include"],
        ["Meridian Corp", 30000, 50, "Include"],
        ["Vanguard Apparel", 20000, 75, "Include"],
        ["Dispute Holdings", 45000, 0, "Do Not Include"],
    ]: ws.append(row)
    for r in range(2, 6): ws.cell(r, 2).number_format = '$#,##0;[Red]-$#,##0'; ws.cell(r, 3).number_format = '0'
    style_sheet(ws, "D", 5)
    for col, width in {"A":24,"B":16,"C":24,"D":20}.items(): ws.column_dimensions[col].width = width

    ws = wb.create_sheet("AP Due")
    ws.append(["Vendor", "Amount", "Status"])
    for row in [
        ["TechServe", 22000, "Include"],
        ["Payroll", 85000, "Include"],
        ["CloudHost Pro", 9600, "Include"],
        ["Cancelled Lease", 18000, "Do Not Include"],
    ]: ws.append(row)
    for r in range(2, 6): ws.cell(r, 2).number_format = '$#,##0;[Red]-$#,##0'
    style_sheet(ws, "C", 5)
    for col, width in {"A":24,"B":16,"C":20}.items(): ws.column_dimensions[col].width = width

    out = ROOT / "data/input_files/cash_flow_inputs.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    wb.save(ROOT / "environment/input_files/cash_flow_inputs.xlsx")

if __name__ == "__main__": main()
