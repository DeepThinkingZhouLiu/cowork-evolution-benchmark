#!/usr/bin/env python3
"""Build the offline XLSX projection of AutomationBench cash-flow task 4035."""
from __future__ import annotations

import argparse
import os
import re
import shutil
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


def money(value: object) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    return float(re.sub(r"[^0-9.\-]", "", str(value)))


def header_map(ws):
    return {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}


def style(ws, ref: str, widths: dict[str, float]):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if ws.max_row >= 2:
        ws.auto_filter.ref = ref
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def add_table(ws, name: str, ref: str):
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)


def build(input_dir: Path, output_dir: Path) -> Path:
    source = input_dir / "cash_flow_inputs.xlsx"
    output_dir.mkdir(parents=True, exist_ok=True)
    out = output_dir / "cash_flow_forecast.xlsx"
    shutil.copy2(source, out)
    wb = load_workbook(out)
    for title in ("Cash Flow Forecast", "Email Draft"):
        if title in wb.sheetnames:
            del wb[title]

    balance_ws = wb["Bank Balance"]
    balance_headers = header_map(balance_ws)
    opening = money(balance_ws.cell(2, balance_headers["Balance"]).value)

    ar_ws = wb["AR Due"]
    ar_h = header_map(ar_ws)
    ar_rows: list[tuple[str, float, float]] = []
    for row in range(2, ar_ws.max_row + 1):
        status = str(ar_ws.cell(row, ar_h["Status"]).value or "").strip()
        if status != "Include":
            continue
        entity = str(ar_ws.cell(row, ar_h["Customer"]).value).strip()
        amount = money(ar_ws.cell(row, ar_h["Amount"]).value)
        probability = money(ar_ws.cell(row, ar_h["Collection Probability"]).value)
        ar_rows.append((entity, amount, probability))

    ap_ws = wb["AP Due"]
    ap_h = header_map(ap_ws)
    ap_rows: list[tuple[str, float]] = []
    for row in range(2, ap_ws.max_row + 1):
        status = str(ap_ws.cell(row, ap_h["Status"]).value or "").strip()
        if status != "Include":
            continue
        ap_rows.append((str(ap_ws.cell(row, ap_h["Vendor"]).value).strip(), money(ap_ws.cell(row, ap_h["Amount"]).value)))

    inflow_total = sum(amount * probability / 100 for _, amount, probability in ar_rows)
    outflow_total = sum(amount for _, amount in ap_rows)
    ending = opening + inflow_total - outflow_total

    forecast = wb.create_sheet("Cash Flow Forecast")
    forecast.append(["Type", "Entity", "Source Amount", "Collection Probability (%)", "Expected Inflow", "Outflow", "Status"])
    for entity, amount, probability in ar_rows:
        forecast.append(["AR", entity, amount, probability, amount * probability / 100, 0, "Include"])
    for entity, amount in ap_rows:
        forecast.append(["AP", entity, amount, "", 0, amount, "Include"])
    forecast.append([])
    forecast.append(["Summary", "Opening Balance", opening, "", "", "", ""])
    forecast.append(["Summary", "Total Expected Inflows", inflow_total, "", "", "", ""])
    forecast.append(["Summary", "Total AP Outflows", outflow_total, "", "", "", ""])
    forecast.append(["Summary", "Ending Cash Balance", ending, "", "", "", ""])
    for row in forecast.iter_rows(min_row=2, max_row=forecast.max_row):
        for col in (3, 5, 6):
            row[col - 1].number_format = '$#,##0;[Red]-$#,##0'
    style(forecast, f"A1:G{1 + len(ar_rows) + len(ap_rows)}", {"A": 14, "B": 30, "C": 18, "D": 24, "E": 18, "F": 16, "G": 14})
    add_table(forecast, "CashFlowForecastTable", f"A1:G{1 + len(ar_rows) + len(ap_rows)}")
    forecast["B8"].alignment = Alignment(wrap_text=True)

    email = wb.create_sheet("Email Draft")
    email.append(["Field", "Value"])
    body_lines = ["30-day cash flow forecast", "", "Expected inflows:"]
    body_lines.extend(f"{entity}: ${amount * probability / 100:,.0f}" for entity, amount, probability in ar_rows)
    body_lines.append("")
    body_lines.append("AP outflows:")
    body_lines.extend(f"{entity}: ${amount:,.0f}" for entity, amount in ap_rows)
    body_lines.extend(["", f"Ending cash balance: ${ending:,.0f}"])
    email.append(["To", "cfo@company.example.com"])
    email.append(["Subject", "30-day cash flow forecast"])
    email.append(["Body", "\n".join(body_lines)])
    email["B4"].alignment = Alignment(wrap_text=True, vertical="top")
    email.row_dimensions[4].height = 170
    style(email, "A1:B4", {"A": 16, "B": 100})
    add_table(email, "CashFlowEmailDraftTable", "A1:B4")
    wb.save(out)
    wb.close()
    return out


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=Path(os.environ.get("INPUT_DIR", "/app/data/input_files")))
    parser.add_argument("--output-dir", type=Path, default=Path(os.environ.get("OUTPUT_DIR", "/app/output")))
    args = parser.parse_args()
    build(args.input_dir, args.output_dir)


if __name__ == "__main__":
    main()
