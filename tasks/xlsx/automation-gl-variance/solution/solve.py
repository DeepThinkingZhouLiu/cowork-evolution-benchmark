#!/usr/bin/env python3
"""AutomationBench finance.gl_variance_analysis 的离线 XLSX 参考解法。"""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


INPUT_NAME = "gl_variance_source.xlsx"
OUTPUT_NAME = "GL_Variance_Analysis.xlsx"
HEADERS = [
    "Account",
    "Type",
    "January_Actual",
    "Comparison_Period",
    "Comparison_Amount",
    "Variance_Amount",
    "Variance_Percentage",
    "Materiality_Basis",
    "Explanation",
]


def money(value: float) -> str:
    return f"${value:,.2f}"


def pct(value: float) -> str:
    return f"{value:.2f}%"


def material_rows(input_path: Path) -> list[dict]:
    workbook = load_workbook(input_path, data_only=True)
    actuals = workbook["January Actuals"]
    policy_sheet = workbook["Materiality Policy"]
    policy: dict[str, dict] = {}
    for row in policy_sheet.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        policy[str(row[0])] = {
            "percent": float(row[1]),
            "dollar": None if row[2] in (None, "") else float(row[2]),
        }
    rows: list[dict] = []
    for row in actuals.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        account, account_type, jan, dec, budget = row
        rule = policy[str(account_type)]
        comparisons = [("Dec 2025", dec)]
        if str(account_type).lower() != "balance sheet":
            comparisons.append(("Budget", budget))
        for period, comparison in comparisons:
            jan_value = float(jan)
            comparison_value = float(comparison)
            variance = jan_value - comparison_value
            variance_pct = variance / comparison_value * 100
            percent_hit = abs(variance_pct) > rule["percent"]
            dollar_hit = rule["dollar"] is not None and abs(variance) > rule["dollar"]
            if not (percent_hit or dollar_hit):
                continue
            basis_parts = []
            if percent_hit:
                basis_parts.append(f"|variance %| > {rule['percent']:.0f}%")
            if dollar_hit:
                basis_parts.append(f"|variance| > {money(rule['dollar'])}")
            rows.append(
                {
                    "account": str(account),
                    "type": str(account_type),
                    "jan": jan_value,
                    "period": period,
                    "comparison": comparison_value,
                    "variance": variance,
                    "percentage": variance_pct,
                    "basis": " or ".join(basis_parts),
                    "explanation": f"{period}: {money(jan_value)} vs {money(comparison_value)}; difference {money(variance)} ({pct(variance_pct)}).",
                }
            )
    workbook.close()
    return rows


def style_sheet(sheet, widths: dict[str, float]) -> None:
    navy = "17365D"
    white = "FFFFFF"
    thin = Side(style="thin", color="D9E2F3")
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 28
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=10, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows():
        for cell in row:
            cell.border = Border(bottom=thin)
            if cell.row > 1:
                cell.font = Font(name="Aptos", size=10, color="1F2937")
                cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (8, 9))
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0


def add_table(sheet, name: str) -> None:
    table = Table(displayName=name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)


def write_workbook(output_path: Path, rows: list[dict]) -> None:
    workbook = Workbook()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.properties.creator = "Cowork Evolution Benchmark"
    workbook.properties.title = "GL Variance Analysis"
    analysis = workbook.active
    analysis.title = "Material Variances"
    analysis.append(HEADERS)
    for row in rows:
        analysis.append([
            row["account"], row["type"], row["jan"], row["period"], row["comparison"],
            row["variance"], row["percentage"], row["basis"], row["explanation"],
        ])
    for row in analysis.iter_rows(min_row=2):
        for index in (3, 5, 6):
            row[index - 1].number_format = "$#,##0.00;[Red]-$#,##0.00"
        row[6].number_format = '0.00"%";[Red]-0.00"%"'
    style_sheet(analysis, {"A": 28, "B": 14, "C": 18, "D": 18, "E": 20, "F": 18, "G": 20, "H": 28, "I": 66})
    add_table(analysis, "MaterialVariances")

    draft = workbook.create_sheet("Controller Email Draft")
    draft.append(["Field", "Value"])
    body_lines = ["The January 2026 GL variance analysis is complete. Material variances only:"]
    for row in rows:
        body_lines.append(
            f"- {row['account']} ({row['type']}) vs {row['period']}: "
            f"Jan {money(row['jan'])}; comparison {money(row['comparison'])}; "
            f"variance {money(row['variance'])} ({pct(row['percentage'])})."
        )
    body_lines.append("Non-material comparisons are omitted per the controller's policy.")
    body = "\n".join(body_lines)
    draft.append(["From", "accounting@company.example.com"])
    draft.append(["To", "controller@company.example.com"])
    draft.append(["Subject", "January 2026 GL Variance Analysis"])
    draft.append(["Body", body])
    draft.append(["Status", "Draft - Not Sent"])
    style_sheet(draft, {"A": 18, "B": 110})
    draft["B5"].alignment = Alignment(vertical="top", wrap_text=True)
    draft.row_dimensions[5].height = max(70, 18 * (body.count("\n") + 2))
    add_table(draft, "ControllerEmailDraft")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=Path("/app/data/input_files"))
    parser.add_argument("--output-dir", type=Path, default=Path("/app/output"))
    args = parser.parse_args()
    write_workbook(args.output_dir / OUTPUT_NAME, material_rows(args.input_dir / INPUT_NAME))


if __name__ == "__main__":
    main()
