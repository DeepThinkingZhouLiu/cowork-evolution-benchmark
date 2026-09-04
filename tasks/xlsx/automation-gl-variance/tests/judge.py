#!/usr/bin/env python3
"""从冻结的 GL 快照和控制器策略独立重算 AutomationBench 差异结果。"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook


TASK_ID = "cowork-evo/automation-gl-variance"
RUBRIC_VERSION = "automation-gl-variance-static-v0.1"
JUDGE_VERSION = "automation-gl-variance-deterministic-v0.1"
OUTPUT_NAME = "GL_Variance_Analysis.xlsx"
INPUT_NAME = "gl_variance_source.xlsx"
HEADERS = [
    "Account", "Type", "January_Actual", "Comparison_Period", "Comparison_Amount",
    "Variance_Amount", "Variance_Percentage", "Materiality_Basis", "Explanation",
]
WEIGHTS = {"R001": 0.10, "R002": 0.25, "R003": 0.22, "R004": 0.13, "R005": 0.15, "R006": 0.10, "R007": 0.05, "R008": -0.10}
EXPECTED_HASH = "3d437aa85af640fa0533095764d8cffc07f70186f2d7c39d99f3980210e64b67"


def digest(path: Path) -> str:
    value = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            value.update(block)
    return value.hexdigest()


def criterion(cid: str, score: bool | float, evidence: str) -> dict:
    return {"criterion_id": cid, "score": float(score), "evidence": evidence, "method": "deterministic"}


def derive_expected(input_path: Path) -> list[dict]:
    workbook = load_workbook(input_path, data_only=True)
    policy_sheet = workbook["Materiality Policy"]
    policy = {}
    for row in policy_sheet.iter_rows(min_row=2, values_only=True):
        if row[0]:
            policy[str(row[0])] = (float(row[1]), None if row[2] in (None, "") else float(row[2]))
    expected = []
    for row in workbook["January Actuals"].iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        account, account_type, jan, dec, budget = row
        pct_threshold, dollar_threshold = policy[str(account_type)]
        comparisons = [("Dec 2025", dec)]
        if str(account_type).lower() != "balance sheet":
            comparisons.append(("Budget", budget))
        for period, comparison in comparisons:
            jan_value = float(jan)
            comparison_value = float(comparison)
            variance = jan_value - comparison_value
            percentage = variance / comparison_value * 100
            percent_hit = abs(percentage) > pct_threshold
            dollar_hit = dollar_threshold is not None and abs(variance) > dollar_threshold
            if percent_hit or dollar_hit:
                expected.append({"account": str(account), "type": str(account_type), "jan": jan_value, "period": period, "comparison": comparison_value, "variance": variance, "percentage": percentage})
    workbook.close()
    return expected


def pct_value(value, number_format: str) -> float:
    number = float(value)
    # 0.00% is the Excel fractional representation; 0.00"%" stores percent points.
    if "%" in number_format and '"%"' not in number_format and abs(number) <= 1.0:
        return number * 100
    return number


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    root = Path(__file__).parents[1]
    input_path = root / "data/input_files" / INPUT_NAME
    output_path = args.artifact_dir / OUTPUT_NAME
    results: list[dict] = []
    expected = []
    try:
        expected = derive_expected(input_path)
        valid_file = output_path.is_file() and zipfile.is_zipfile(output_path)
        if not valid_file:
            results.append(criterion("R001", False, "输出缺失或不是有效 OOXML 工作簿"))
            for cid in ("R002", "R003", "R004", "R005"):
                results.append(criterion(cid, False, "缺少有效输出"))
            results.append(criterion("R006", False, "缺少有效输出"))
        else:
            workbook = load_workbook(output_path, data_only=False)
            structure = workbook.sheetnames == ["Material Variances", "Controller Email Draft"]
            results.append(criterion("R001", structure, f"工作表={workbook.sheetnames}"))
            analysis = workbook["Material Variances"]
            headers = [str(cell.value or "").strip() for cell in analysis[1]]
            rows = []
            for row in analysis.iter_rows(min_row=2, values_only=False):
                if row[0].value is None:
                    continue
                rows.append({
                    "account": str(row[0].value).strip(),
                    "type": str(row[1].value).strip(),
                    "jan": float(row[2].value),
                    "period": str(row[3].value).strip(),
                    "comparison": float(row[4].value),
                    "variance": float(row[5].value),
                    "percentage": pct_value(row[6].value, row[6].number_format),
                    "basis": str(row[7].value or ""),
                    "explanation": str(row[8].value or ""),
                })
            core_headers = headers == HEADERS
            identity_ok = len(rows) == len(expected) and all(
                out["account"] == exp["account"] and out["type"] == exp["type"] and out["period"] == exp["period"]
                for out, exp in zip(rows, expected)
            )
            results.append(criterion("R002", core_headers and identity_ok, f"表头={core_headers}; 行数={len(rows)}; 期望={len(expected)}"))
            numeric_ok = len(rows) == len(expected) and all(
                abs(out[key] - exp[key]) <= (0.02 if key == "percentage" else 0.01)
                for out, exp in zip(rows, expected)
                for key in ("jan", "comparison", "variance", "percentage")
            )
            results.append(criterion("R003", numeric_ok, "实际数、比较数、差异金额和百分比独立重算"))
            excluded = {"Revenue - Consulting", "Salaries Expense", "Office Supplies"}
            strict_ok = identity_ok and not any(row["account"] in excluded for row in rows) and len(rows) == 3
            results.append(criterion("R004", strict_ok, "严格阈值和非重大账户排除检查"))
            draft = workbook["Controller Email Draft"]
            fields = {str(row[0].value).strip(): row[1].value for row in draft.iter_rows(min_row=2, max_col=2) if row[0].value}
            body = str(fields.get("Body", ""))
            draft_ok = (
                fields.get("To") == "controller@company.example.com"
                and "GL Variance Analysis" in str(fields.get("Subject", ""))
                and fields.get("Status") == "Draft - Not Sent"
                and all(token in body for token in ("Revenue - SaaS", "25,000", "Marketing Expense", "13,000", "8,000"))
                and all(token not in body for token in ("Consulting", "Office Supplies"))
            )
            results.append(criterion("R005", draft_ok, "草稿收件人、主题、正文重大差异和未发送状态检查"))
            professional = all(
                workbook[name].freeze_panes == "A2" and bool(workbook[name].tables) and workbook[name].auto_filter.ref
                for name in workbook.sheetnames
            ) and all((analysis.column_dimensions[c].width or 0) >= 12 for c in "ABCDEFGHI")
            results.append(criterion("R006", professional, "原生表格、筛选、冻结和列宽检查"))
            workbook.close()
        data_input = root / "data/input_files" / INPUT_NAME
        env_input = root / "environment/input_files" / INPUT_NAME
        trace_ok = digest(data_input) == EXPECTED_HASH and env_input.read_bytes() == data_input.read_bytes()
        results.append(criterion("R007", trace_ok, f"输入哈希={digest(data_input)}; data/environment 字节一致={env_input.read_bytes() == data_input.read_bytes()}"))
        violation = False
        if output_path.is_file() and zipfile.is_zipfile(output_path):
            raw = output_path.read_bytes()
            violation = b"sent successfully" in raw.lower() or b"consulting" in raw.lower() and len(expected) == 3
        results.append(criterion("R008", 1.0 if violation else 0.0, "完整性惩罚项；未检测到伪造发送或遗漏约束" if not violation else "检测到不应出现的发送声明或排除账户"))
    except Exception as exc:
        results = [criterion("R001", False, f"Judge 读取失败：{exc}")]
        for cid in ("R002", "R003", "R004", "R005", "R006", "R007", "R008"):
            results.append(criterion(cid, False, "读取失败"))
    values = {item["criterion_id"]: item["score"] for item in results}
    reward = sum(WEIGHTS[cid] * values.get(cid, 0.0) for cid in WEIGHTS)
    passed = reward >= 0.8 and all(values.get(cid, 0.0) == 1.0 for cid in ("R001", "R002", "R003", "R004", "R005"))
    payload = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [digest(output_path)] if output_path.is_file() else [], "criterion_results": results, "reward": reward, "passed": passed}
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.write_text(f"{reward:.6f}\n", encoding="utf-8")
    raise SystemExit(0 if passed else 1)


if __name__ == "__main__":
    main()
