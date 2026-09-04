"""Create the solver-visible workbook from the fixed AutomationBench fixture."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).parents[1]

def style_sheet(ws, end_col, end_row):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{end_col}{end_row}"
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    tab = Table(displayName=f"{ws.title.replace(' ', '')}Table", ref=f"A1:{end_col}{end_row}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)

def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Vendor Statements Received"
    ws.append(["Vendor", "Email", "Statement Balance"])
    for row in [
        ["Acme Supplies", "ar@acme.example.com", 6200],
        ["CloudHost Pro", "billing@cloudhost.example.com", 9600],
        ["Metro Supply", "ar@metro.example.com", 1800],
        ["TechServe Solutions", "billing@techserve.example.com", 4500],
    ]: ws.append(row)
    for r in range(2, 6): ws.cell(r, 3).number_format = '$#,##0;[Red]-$#,##0'
    style_sheet(ws, "C", 5)
    for col, width in {"A":24,"B":34,"C":20}.items(): ws.column_dimensions[col].width = width

    ws = wb.create_sheet("Our AP Balances")
    ws.append(["Vendor", "AP Balance", "Notes"])
    for row in [
        ["Acme Supplies", 4200, ""],
        ["CloudHost Pro", 9600, ""],
        ["Metro Supply", 2580, ""],
        ["TechServe Solutions", 4500, "payment in transit - exclude from recon"],
    ]: ws.append(row)
    for r in range(2, 6): ws.cell(r, 2).number_format = '$#,##0;[Red]-$#,##0'
    style_sheet(ws, "C", 5)
    for col, width in {"A":24,"B":16,"C":42}.items(): ws.column_dimensions[col].width = width

    out = ROOT / "data/input_files/vendor_reconciliation_inputs.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    wb.save(ROOT / "environment/input_files/vendor_reconciliation_inputs.xlsx")

if __name__ == "__main__": main()
