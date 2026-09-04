#!/usr/bin/env python3
"""Build the offline XLSX projection of AutomationBench vendor reconciliation 4075."""
from __future__ import annotations

import argparse
import os
import re
import shutil
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


def money(value: object) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    return float(re.sub(r"[^0-9.\-]", "", str(value)))


def header_map(ws):
    return {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}


def style(ws, ref: str, widths: dict[str, float]):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.auto_filter.ref = ref
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def add_table(ws, name: str, ref: str):
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)


def build(input_dir: Path, output_dir: Path) -> Path:
    source = input_dir / "vendor_reconciliation_inputs.xlsx"
    output_dir.mkdir(parents=True, exist_ok=True)
    out = output_dir / "vendor_reconciliation.xlsx"
    shutil.copy2(source, out)
    wb = load_workbook(out)
    for title in ("Reconciliation", "Email Draft"):
        if title in wb.sheetnames:
            del wb[title]

    statements = wb["Vendor Statements Received"]
    sh = header_map(statements)
    statement_rows: list[tuple[str, str, float]] = []
    for row in range(2, statements.max_row + 1):
        statement_rows.append((str(statements.cell(row, sh["Vendor"]).value).strip(), str(statements.cell(row, sh["Email"]).value).strip(), money(statements.cell(row, sh["Statement Balance"]).value)))

    ap = wb["Our AP Balances"]
    ah = header_map(ap)
    ap_rows: dict[str, tuple[float, str]] = {}
    for row in range(2, ap.max_row + 1):
        ap_rows[str(ap.cell(row, ah["Vendor"]).value).strip()] = (money(ap.cell(row, ah["AP Balance"]).value), str(ap.cell(row, ah["Notes"]).value or "").strip())

    recon = wb.create_sheet("Reconciliation")
    recon.append(["Vendor", "Email", "Statement Balance", "AP Balance", "Difference", "Status", "Contact Required", "Notes"])
    unresolved: list[tuple[str, str, float, float, float]] = []
    for vendor, email, statement_balance in statement_rows:
        ap_balance, notes = ap_rows[vendor]
        excluded = "exclude from recon" in notes.lower()
        difference = 0.0 if excluded else statement_balance - ap_balance
        if excluded:
            status, contact, note = "Excluded", "No", notes
        elif abs(difference) < 1e-9:
            status, contact, note = "Matched", "No", ""
        else:
            status, contact, note = "Unresolved discrepancy", "Yes", "Statement less AP balance"
            unresolved.append((vendor, email, statement_balance, ap_balance, difference))
        recon.append([vendor, email, statement_balance, ap_balance, difference, status, contact, note])
    for row in recon.iter_rows(min_row=2, max_row=recon.max_row):
        for col in (3, 4, 5): row[col - 1].number_format = '$#,##0;[Red]-$#,##0'
    style(recon, f"A1:H{recon.max_row}", {"A":24, "B":34, "C":20, "D":16, "E":16, "F":24, "G":18, "H":40})
    add_table(recon, "VendorReconciliationTable", f"A1:H{recon.max_row}")

    email_sheet = wb.create_sheet("Email Draft")
    email_sheet.append(["Recipient", "Subject", "Body"])
    summary_lines = ["Vendor statement reconciliation summary", "", "Unresolved discrepancies:"]
    for vendor, _, statement_balance, ap_balance, difference in unresolved:
        summary_lines.append(f"{vendor}: statement ${statement_balance:,.0f}; AP ${ap_balance:,.0f}; difference ${difference:,.0f}")
    email_sheet.append(["controller@company.example.com", "Vendor reconciliation summary", "\n".join(summary_lines)])
    for vendor, address, statement_balance, ap_balance, difference in unresolved:
        body = "\n".join([
            f"Hello {vendor},",
            "Our reconciliation identified an unresolved balance difference.",
            f"Statement balance: ${statement_balance:,.0f}",
            f"Our AP balance: ${ap_balance:,.0f}",
            f"Difference: ${difference:,.0f}",
        ])
        email_sheet.append([address, f"Reconciliation discrepancy - {vendor}", body])
    for row in range(2, email_sheet.max_row + 1):
        email_sheet.cell(row, 3).alignment = Alignment(wrap_text=True, vertical="top")
        email_sheet.row_dimensions[row].height = 110
    style(email_sheet, f"A1:C{email_sheet.max_row}", {"A": 34, "B": 42, "C": 90})
    add_table(email_sheet, "VendorEmailDraftTable", f"A1:C{email_sheet.max_row}")
    wb.save(out)
    wb.close()
    return out


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=Path(os.environ.get("INPUT_DIR", "/app/data/input_files")))
    parser.add_argument("--output-dir", type=Path, default=Path(os.environ.get("OUTPUT_DIR", "/app/output")))
    args = parser.parse_args()
    build(args.input_dir, args.output_dir)


if __name__ == "__main__":
    main()
