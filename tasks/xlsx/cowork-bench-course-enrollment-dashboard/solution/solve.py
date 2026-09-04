#!/usr/bin/env python3
"""将 Canvas 离线快照计算为一个可重算的课程风险 XLSX。"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


RISK_ORDER = ["Critical", "At Risk", "On Track"]
COURSE_IDS = [16, 17]
NAVY = "17365D"
BLUE = "2F75B5"
WHITE = "FFFFFF"
INK = "1F2937"
PALE = "EAF2F8"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream))


def number(value: str | None) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        return float(str(value).strip().replace(",", "."))
    except ValueError:
        return None


def truthy(value: str | None) -> bool:
    return str(value or "").strip().casefold() in {"t", "true", "1", "yes", "y"}


def parse_thresholds(model: dict) -> tuple[float, float]:
    high_text = model["risk_classification"]["high_risk"]["threshold"]
    medium_text = model["risk_classification"]["medium_risk"]["threshold"]
    high = re.findall(r"-?\d+(?:\.\d+)?", high_text)
    medium = re.findall(r"-?\d+(?:\.\d+)?", medium_text)
    if not high or len(medium) < 2:
        raise ValueError("scoring_model.json 中缺少连续风险阈值")
    high_cutoff = float(high[0])
    low_cutoff = float(medium[1]) + 0.01
    if float(medium[0]) != high_cutoff:
        raise ValueError("风险阈值不连续")
    return high_cutoff, round(low_cutoff, 2)


def classify(value: float, high_cutoff: float, low_cutoff: float) -> str:
    if value < high_cutoff:
        return "Critical"
    if value < low_cutoff:
        return "At Risk"
    return "On Track"


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def derive(input_dir: Path) -> dict:
    courses = {
        int(row["course_id"]): row
        for row in read_csv(input_dir / "canvas_courses.csv")
        if int(row["course_id"]) in COURSE_IDS
    }
    if sorted(courses) != COURSE_IDS:
        raise ValueError(f"课程快照必须包含 {COURSE_IDS}，实际为 {sorted(courses)}")
    enrollments = read_csv(input_dir / "canvas_enrollments.csv")
    assignments = read_csv(input_dir / "canvas_assignments.csv")
    assignment_course = {
        int(row["assignment_id"]): int(row["course_id"])
        for row in assignments
        if int(row["course_id"]) in COURSE_IDS
    }
    if len(assignment_course) != 26:
        raise ValueError(f"课程 16/17 应有 26 个作业，实际 {len(assignment_course)}")
    model = json.loads((input_dir / "scoring_model.json").read_text(encoding="utf-8"))
    high_cutoff, low_cutoff = parse_thresholds(model)

    enrollment_by_course: dict[int, list[dict[str, str]]] = defaultdict(list)
    for row in enrollments:
        cid = int(row["course_id"])
        if cid in COURSE_IDS:
            enrollment_by_course[cid].append(row)
    student_score_by: dict[tuple[int, str], float] = {}
    score_values: dict[int, list[float]] = defaultdict(list)
    for row in enrollments:
        cid = int(row["course_id"])
        score = number(row.get("current_score"))
        if cid in COURSE_IDS and "student" in row.get("type", "").casefold() and score is not None:
            student_score_by[(cid, str(row["user_id"]))] = score
            score_values[cid].append(score)

    submission_counts: Counter[int] = Counter()
    graded_submission_counts: Counter[int] = Counter()
    late_submission_counts: Counter[int] = Counter()
    for row in read_csv(input_dir / "canvas_submissions.csv"):
        assignment_id = int(row["assignment_id"])
        cid = assignment_course.get(assignment_id)
        if cid is None:
            continue
        submission_counts[cid] += 1
        if row.get("workflow_state", "").strip().casefold() == "graded":
            graded_submission_counts[cid] += 1
        if truthy(row.get("late")):
            late_submission_counts[cid] += 1

    metrics = []
    for cid in COURSE_IDS:
        rows = enrollment_by_course[cid]
        scores = score_values[cid]
        if not scores:
            continue
        graded_students = sum(1 for key in student_score_by if key[0] == cid)
        total_submissions = submission_counts[cid]
        metrics.append({
            "Course_ID": cid,
            "Course_Name": courses[cid]["course_name"],
            "Course_Code": courses[cid]["course_code"],
            "Total_Enrollment": len(rows),
            "Student_Enrollment": sum("student" in row.get("type", "").casefold() for row in rows),
            "Graded_Students": graded_students,
            "Submission_Count": total_submissions,
            "Graded_Submissions": graded_submission_counts[cid],
            "Avg_Score": round(sum(scores) / len(scores), 2),
            "Completion_Rate": round(100.0 * graded_submission_counts[cid] / total_submissions, 1) if total_submissions else 0.0,
            "Late_Rate": round(100.0 * late_submission_counts[cid] / total_submissions, 1) if total_submissions else 0.0,
            "Late_Submissions": late_submission_counts[cid],
        })
    metrics.sort(key=lambda row: (row["Avg_Score"], row["Course_ID"]))
    for row in metrics:
        row["Risk_Level"] = classify(row["Avg_Score"], high_cutoff, low_cutoff)

    distribution = []
    for risk in RISK_ORDER:
        members = [row for row in metrics if row["Risk_Level"] == risk]
        distribution.append({
            "Risk_Level": risk,
            "Course_Count": len(members),
            "Student_Enrollment": sum(row["Student_Enrollment"] for row in members),
        })
    actions = model["intervention_actions"]
    action_rows = []
    for risk, key in [("Critical", "high"), ("At Risk", "medium"), ("On Track", "low")]:
        action_rows.append({"Risk_Level": risk, **actions[key]})
    return {
        "metrics": metrics,
        "distribution": distribution,
        "actions": action_rows,
        "thresholds": {"critical_below": high_cutoff, "at_risk_below": low_cutoff},
        "input_hashes": {p.name: digest(p) for p in sorted(input_dir.iterdir()) if p.is_file()},
    }


def style_sheet(ws, widths: dict[str, float]) -> None:
    header_fill = PatternFill("solid", fgColor=NAVY)
    thin = Side(style="thin", color="D7E0EA")
    for cell in ws[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color=INK)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35


def add_table(ws, ref: str, name: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)


def write_workbook(path: Path, data: dict) -> None:
    wb = Workbook()
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    metrics = data["metrics"]

    ws = wb.active
    ws.title = "Course_Metrics"
    headers = ["Course_ID", "Course_Name", "Course_Code", "Total_Enrollment", "Student_Enrollment", "Graded_Students", "Submission_Count", "Graded_Submissions", "Avg_Score", "Completion_Rate", "Late_Rate", "Late_Submissions", "Risk_Level"]
    ws.append(headers)
    for row in metrics:
        ws.append([row[h] for h in headers])
    style_sheet(ws, {"A": 12, "B": 40, "C": 16, "D": 18, "E": 19, "F": 17, "G": 18, "H": 20, "I": 13, "J": 18, "K": 13, "L": 17, "M": 14})
    for cell in ws["I"][1:]: cell.number_format = "0.00"
    for col in ("J", "K"):
        for cell in ws[col][1:]: cell.number_format = '0.0"%"'
    add_table(ws, f"A1:L{len(metrics)+1}", "CourseMetrics")

    ws = wb.create_sheet("Risk_Distribution")
    ws.append(["Risk_Level", "Course_Count", "Student_Enrollment"])
    for idx, row in enumerate(data["distribution"], start=2):
        ws.append([row["Risk_Level"], None, None])
        ws[f"B{idx}"] = f'=COUNTIF(Course_Metrics!$M$2:$M${len(metrics)+1},A{idx})'
        ws[f"C{idx}"] = f'=SUMIF(Course_Metrics!$M$2:$M${len(metrics)+1},A{idx},Course_Metrics!$E$2:$E${len(metrics)+1})'
    style_sheet(ws, {"A": 16, "B": 17, "C": 21})
    add_table(ws, "A1:C4", "RiskDistribution")

    ws = wb.create_sheet("Submission_Quality")
    q_headers = ["Course_Name", "Submission_Count", "Graded_Submissions", "Completion_Rate", "Late_Submissions", "Late_Rate"]
    ws.append(q_headers)
    for idx, row in enumerate(metrics, start=2):
        source_row = idx
        ws.append([
            f"=Course_Metrics!B{source_row}",
            f"=Course_Metrics!G{source_row}",
            f"=Course_Metrics!H{source_row}",
            f"=Course_Metrics!J{source_row}",
            f"=Course_Metrics!L{source_row}",
            f"=Course_Metrics!K{source_row}",
        ])
    style_sheet(ws, {"A": 40, "B": 18, "C": 20, "D": 18, "E": 17, "F": 13})
    for col in ("D", "F"):
        for cell in ws[col][1:]: cell.number_format = '0.0"%"'
    add_table(ws, f"A1:F{len(metrics)+1}", "SubmissionQuality")

    ws = wb.create_sheet("Risk_Action_Plan")
    ws.append(["Risk_Level", "Action", "Timeline", "Responsible"])
    for row in data["actions"]:
        ws.append([row["Risk_Level"], row["action"], row["timeline"], row["responsible"]])
    style_sheet(ws, {"A": 16, "B": 38, "C": 18, "D": 26})
    add_table(ws, "A1:D4", "RiskActionPlan")

    ws = wb.create_sheet("Source_Notes")
    ws.append(["Item", "Details"])
    notes = [
        ("Source", "Cowork-Bench course-enrollment-analytics-dashboard; Canvas rows projected from fixed db/init.sql.gz revision d943e75bc0fc8e3b27141979300cd8cbcd1e890d"),
        ("Risk thresholds", "Critical: Avg_Score < 60; At Risk: 60 <= Avg_Score < 75; On Track: Avg_Score >= 75"),
        ("Score basis", "Avg_Score uses non-empty current_score from StudentEnrollment rows; no scores were invented"),
        ("Submission metrics", "Completion_Rate = graded submissions / all submissions; Late_Rate = late=true submissions / all submissions"),
        ("Inputs", "canvas_courses.csv, canvas_enrollments.csv, canvas_assignments.csv, canvas_submissions.csv, scoring_model.json, Retention_Policy.pdf"),
        ("Input SHA256", "; ".join(f"{name}={value}" for name, value in sorted(data["input_hashes"].items()))),
        ("Status", "author_hypothesis=L4; calibration_status=uncalibrated; review_status=pending_human_review"),
    ]
    for row in notes: ws.append(list(row))
    style_sheet(ws, {"A": 22, "B": 120})
    add_table(ws, f"A1:B{len(notes)+1}", "SourceNotes")

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=Path("/app/data/input_files"))
    parser.add_argument("--output-dir", type=Path, default=Path("/app/output"))
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    data = derive(args.input_dir)
    write_workbook(args.output_dir / "Course_Enrollment_Analytics.xlsx", data)


if __name__ == "__main__":
    main()
