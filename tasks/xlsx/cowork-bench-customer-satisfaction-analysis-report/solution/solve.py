#!/usr/bin/env python3
"""Cowork-Bench 客户满意度分析到报告题的逐题参考解法。"""

from __future__ import annotations

import argparse
import csv
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


ANALYSIS_NAME = "satisfaction_analysis.xlsx"
ACTION_NAME = "action_plans.xlsx"
REPORT_NAME = "satisfaction_report.docx"
SUMMARY_NAME = "executive_summary.docx"
SEGMENT_ORDER = ["Enterprise", "Mid-Market", "SMB"]


def d(value) -> Decimal:
    return Decimal(str(value))


def round_2(value: Decimal) -> Decimal:
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def average(values) -> Decimal:
    items = [d(value) for value in values]
    return round_2(sum(items, Decimal("0")) / len(items))


def nps(rows: list[dict]) -> Decimal:
    promoters = sum(d(row["Recommend_Score"]) >= 9 for row in rows)
    detractors = sum(d(row["Recommend_Score"]) <= 6 for row in rows)
    return round_2(Decimal(promoters - detractors) * 100 / len(rows))


def calculate(input_dir: Path) -> dict:
    workbook = load_workbook(input_dir / "survey_template.xlsx", data_only=True)
    response_sheet = workbook["Responses"]
    response_rows = list(response_sheet.iter_rows(values_only=True))
    headers = [str(value) for value in response_rows[0]]
    responses = [dict(zip(headers, row)) for row in response_rows[1:] if row and row[0]]
    workbook.close()

    with (input_dir / "support_data_export.csv").open(encoding="utf-8", newline="") as handle:
        tickets = list(csv.DictReader(handle))

    dimensions = {
        "Product Quality": (average(row["Product_Quality"] for row in responses), "1-5"),
        "Customer Service": (average(row["Customer_Service"] for row in responses), "1-5"),
        "Value for Money": (average(row["Value_for_Money"] for row in responses), "1-5"),
        "Ease of Use": (average(row["Ease_of_Use"] for row in responses), "1-5"),
        "Average Recommend Score": (average(row["Recommend_Score"] for row in responses), "0-10"),
    }
    segments = defaultdict(list)
    for row in responses:
        segments[str(row["Customer_Type"])].append(row)
    segment_metrics = {}
    for segment in SEGMENT_ORDER:
        rows = segments[segment]
        segment_metrics[segment] = {
            "count": len(rows),
            "satisfaction": average(row["Overall_Satisfaction"] for row in rows),
            "nps": nps(rows),
            "fcr": round_2(Decimal(sum(str(row["Resolved_First_Contact"]).lower() == "yes" for row in rows)) / len(rows)),
        }

    support_fcr = round_2(Decimal(sum(row["First_Contact_Resolution"].lower() == "yes" for row in tickets)) / len(tickets))
    support_resolution = average(row["Resolution_Time_Hours"] for row in tickets)
    promoters = sum(d(row["Recommend_Score"]) >= 9 for row in responses)
    detractors = sum(d(row["Recommend_Score"]) <= 6 for row in responses)
    return {
        "survey_count": len(responses),
        "support_count": len(tickets),
        "overall_satisfaction": average(row["Overall_Satisfaction"] for row in responses),
        "nps": nps(responses),
        "promoters": promoters,
        "detractors": detractors,
        "support_fcr": support_fcr,
        "average_resolution": support_resolution,
        "dimensions": dimensions,
        "segments": segment_metrics,
    }


def table_style(sheet, ref: str, name: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def section_title(sheet, row: int, text: str, end_column: int) -> None:
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_column)
    cell = sheet.cell(row, 1, text)
    cell.fill = PatternFill("solid", fgColor="D9EAF7")
    cell.font = Font(name="Aptos Display", size=12, bold=True, color="17365D")
    cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[row].height = 24


def apply_sheet_basics(sheet, widths: dict[str, float], freeze: str) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = freeze
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_options.horizontalCentered = True
    sheet.oddFooter.center.text = "Cowork-Bench · Customer Satisfaction Analysis"


def style_header(sheet, row: int, columns: int) -> None:
    for cell in sheet[row][:columns]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 28


def style_body(sheet, min_row: int, max_row: int, columns: int) -> None:
    thin = Side(style="thin", color="D9E2F3")
    for row in sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=1, max_col=columns):
        for cell in row:
            cell.font = Font(name="Aptos", size=10, color="1F2937")
            cell.alignment = Alignment(vertical="center", wrap_text=True, horizontal="left" if cell.column == 1 else "right")
            cell.border = Border(bottom=thin)


def write_analysis(path: Path, metrics: dict) -> None:
    workbook = Workbook()
    workbook.properties.creator = "Cowork Evolution Benchmark"
    workbook.properties.title = "Customer Satisfaction Analysis"
    sheet = workbook.active
    sheet.title = "Satisfaction Analysis"

    section_title(sheet, 1, "Summary Metrics", 3)
    sheet.append(["Metric", "Value", "Source"])
    summary_rows = [
        ("Overall Satisfaction Score", float(metrics["overall_satisfaction"]), f"survey_template.xlsx / Responses (n={metrics['survey_count']})"),
        ("Net Promoter Score (NPS)", float(metrics["nps"]), f"survey_template.xlsx / Responses (n={metrics['survey_count']})"),
        ("First Contact Resolution Rate", float(metrics["support_fcr"]), f"support_data_export.csv (n={metrics['support_count']})"),
        ("Average Resolution Time (Hours)", float(metrics["average_resolution"]), f"support_data_export.csv (n={metrics['support_count']})"),
    ]
    for row in summary_rows:
        sheet.append(row)
    style_header(sheet, 2, 3)
    style_body(sheet, 3, 6, 3)
    sheet["B3"].number_format = "0.00"
    sheet["B4"].number_format = "0.00"
    sheet["B5"].number_format = "0.00%"
    sheet["B6"].number_format = "0.00"
    table_style(sheet, "A2:C6", "Satisfaction_Summary")

    section_title(sheet, 8, "Dimension Ratings", 3)
    for value in ("Dimension", "Average", "Scale"):
        sheet.cell(9, ("Dimension", "Average", "Scale").index(value) + 1, value)
    for dimension, (value, scale) in metrics["dimensions"].items():
        sheet.append([dimension, float(value), scale])
    style_header(sheet, 9, 3)
    style_body(sheet, 10, 14, 3)
    for row in range(10, 15):
        sheet.cell(row, 2).number_format = "0.00"
    table_style(sheet, "A9:C14", "Satisfaction_Dimensions")

    section_title(sheet, 16, "Segment Analysis", 5)
    segment_headers = ["Customer_Type", "Respondents", "Overall_Satisfaction", "NPS", "First_Contact_Resolution_Rate"]
    for column, value in enumerate(segment_headers, start=1):
        sheet.cell(17, column, value)
    for segment in SEGMENT_ORDER:
        item = metrics["segments"][segment]
        sheet.append([segment, item["count"], float(item["satisfaction"]), float(item["nps"]), float(item["fcr"])])
    style_header(sheet, 17, 5)
    style_body(sheet, 18, 20, 5)
    for row in range(18, 21):
        sheet.cell(row, 2).number_format = "0"
        sheet.cell(row, 3).number_format = "0.00"
        sheet.cell(row, 4).number_format = "0.00"
        sheet.cell(row, 5).number_format = "0.00%"
    table_style(sheet, "A17:E20", "Satisfaction_Segments")
    sheet.conditional_formatting.add(
        "C18:C20",
        ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"),
    )
    sheet.conditional_formatting.add(
        "D18:D20",
        ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"),
    )
    apply_sheet_basics(sheet, {"A": 38, "B": 18, "C": 46, "D": 16, "E": 36}, "A3")
    workbook.save(path)


def write_actions(path: Path, metrics: dict) -> list[dict]:
    smb = metrics["segments"]["SMB"]["satisfaction"]
    actions = [
        {
            "Initiative": "Improve First Contact Resolution",
            "Owner": "Head of Support",
            "Target Metric": "First Contact Resolution Rate",
            "Baseline": float(metrics["support_fcr"]),
            "Goal": 0.85,
            "Timeline": "90 days",
            "Status": "Planned",
        },
        {
            "Initiative": "Reduce Resolution Time",
            "Owner": "Support Operations Lead",
            "Target Metric": "Average Resolution Time (Hours)",
            "Baseline": float(metrics["average_resolution"]),
            "Goal": 2.0,
            "Timeline": "120 days",
            "Status": "Planned",
        },
        {
            "Initiative": "SMB Customer Success Program",
            "Owner": "Customer Success Lead",
            "Target Metric": "SMB Overall Satisfaction",
            "Baseline": float(smb),
            "Goal": 4.2,
            "Timeline": "180 days",
            "Status": "Planned",
        },
        {
            "Initiative": "Value Communication Review",
            "Owner": "Product Marketing Lead",
            "Target Metric": "Value for Money",
            "Baseline": float(metrics["dimensions"]["Value for Money"][0]),
            "Goal": 4.2,
            "Timeline": "90 days",
            "Status": "Planned",
        },
    ]
    headers = ["Initiative", "Owner", "Target Metric", "Baseline", "Goal", "Timeline", "Status"]
    workbook = Workbook()
    workbook.properties.creator = "Cowork Evolution Benchmark"
    workbook.properties.title = "Customer Satisfaction Action Plans"
    sheet = workbook.active
    sheet.title = "Action Plans"
    sheet.append(headers)
    for action in actions:
        sheet.append([action[header] for header in headers])
    style_header(sheet, 1, 7)
    style_body(sheet, 2, 5, 7)
    for row in (2,):
        sheet.cell(row, 4).number_format = "0.00%"
        sheet.cell(row, 5).number_format = "0.00%"
    for row in range(3, 6):
        sheet.cell(row, 4).number_format = "0.00"
        sheet.cell(row, 5).number_format = "0.00"
    table_style(sheet, "A1:G5", "Customer_Satisfaction_Actions")
    apply_sheet_basics(sheet, {"A": 38, "B": 28, "C": 38, "D": 16, "E": 16, "F": 16, "G": 16}, "A2")
    workbook.save(path)
    return actions


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_doc_defaults(document: Document, title: str) -> None:
    document.core_properties.title = title
    document.core_properties.author = "Cowork Evolution Benchmark"
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for style_name, size, color in (("Title", 24, "17365D"), ("Heading 1", 16, "17365D"), ("Heading 2", 12, "2E75B6")):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
    footer = section.footer.paragraphs[0]
    footer.text = "Cowork-Bench · Customer Satisfaction"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(100, 116, 139)


def add_metric_table(document: Document, rows: list[tuple[str, str, str]]) -> None:
    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Shading Accent 1"
    for cell, text in zip(table.rows[0].cells, ("Метрика", "Значение", "Источник")):
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade_cell(cell, "17365D")
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for metric, value, source in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, (metric, value, source)):
            cell.text = text
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    table.columns[0].width = Cm(7.0)
    table.columns[1].width = Cm(3.0)
    table.columns[2].width = Cm(7.0)


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.add_run(text)


def write_report(path: Path, metrics: dict, actions: list[dict]) -> None:
    document = Document()
    set_doc_defaults(document, "Анализ удовлетворённости клиентов")
    document.add_heading("Отчёт об удовлетворённости клиентов", 0)
    subtitle = document.add_paragraph("Консолидированный анализ опроса и обращений в службу поддержки")
    subtitle.runs[0].font.color.rgb = RGBColor(71, 85, 105)

    document.add_heading("Резюме", level=1)
    document.add_paragraph(
        f"Анализ охватывает {metrics['survey_count']} ответов клиентов и {metrics['support_count']} обращений в службу поддержки. "
        f"Общая удовлетворённость равна {metrics['overall_satisfaction']:.2f}/5, NPS — {metrics['nps']:.0f}. "
        f"В поддержке {metrics['support_fcr']:.0%} обращений решены с первого контакта, а среднее время решения составляет {metrics['average_resolution']:.2f} часа. "
        "Главный сегментный сигнал — SMB: у него минимальная средняя удовлетворённость 4.00 и NPS -30."
    )

    document.add_heading("Данные и методика", level=1)
    document.add_paragraph(
        "Опрос используется для общей оценки, оценок по измерениям, NPS и сегментной разбивки. "
        "NPS рассчитан как доля респондентов с баллом 9–10 минус доля с баллом 0–6, в пунктах. "
        "Из выгрузки поддержки рассчитаны доля решений с первого контакта и среднее время решения."
    )

    document.add_heading("Ключевые метрики", level=1)
    add_metric_table(
        document,
        [
            ("Overall Satisfaction Score", f"{metrics['overall_satisfaction']:.2f}/5", f"Опрос, n={metrics['survey_count']}"),
            ("Net Promoter Score (NPS)", f"{metrics['nps']:.0f}", f"{metrics['promoters']} промоутеров; {metrics['detractors']} критиков"),
            ("First Contact Resolution Rate", f"{metrics['support_fcr']:.0%}", f"Поддержка, n={metrics['support_count']}"),
            ("Average Resolution Time", f"{metrics['average_resolution']:.2f} часа", f"Поддержка, n={metrics['support_count']}"),
        ],
    )

    document.add_heading("Оценки по измерениям", level=1)
    for label, (value, scale) in metrics["dimensions"].items():
        add_bullet(document, f"{label}: {value:.2f} по шкале {scale}.")
    document.add_paragraph(
        "Лучшая оценка по шкале 1–5 — Ease of Use (4.40), а минимальная — Value for Money (3.90). "
        "Это обосновывает приоритет по разъяснению ценности продукта, не отменяя необходимость улучшать сервис."
    )

    document.add_heading("Сегментный анализ", level=1)
    for segment in SEGMENT_ORDER:
        item = metrics["segments"][segment]
        add_bullet(
            document,
            f"{segment}: {item['count']} ответов; удовлетворённость {item['satisfaction']:.2f}; NPS {item['nps']:.0f}; решение с первого контакта {item['fcr']:.0%}."
        )
    document.add_paragraph(
        "SMB — сегмент с самой низкой средней удовлетворённостью (4.00) и самым низким NPS (-30). "
        "При этом в сегменте всего 20 ответов, поэтому величину разрыва нужно проверить на следующей волне."
    )

    document.add_heading("План улучшений", level=1)
    for action in actions[:3]:
        add_bullet(
            document,
            f"{action['Initiative']} — ответственный {action['Owner']}; срок {action['Timeline']}; метрика {action['Target Metric']}."
        )
    document.add_paragraph(
        "Первый контур — повысить FCR с 75% до 85%; второй — сократить среднее время решения с 3.40 до 2.00 часа; "
        "третий — повысить удовлетворённость SMB с 4.00 до 4.20. Цели являются управленческими ориентирами, а не исходными фактами."
    )

    document.add_heading("Ограничения", level=1)
    document.add_paragraph(
        "В опросе нет Customer_ID, поэтому 50 ответов нельзя построчно связать с 20 обращениями поддержки. "
        "Мы не делаем выводов о корреляции на уровне клиента и не утверждаем причинную связь между скоростью поддержки и оценкой опроса. "
        "Выборки невелики; для решений о масштабировании нужна повторная волна и доверительные интервалы."
    )

    document.add_heading("Вывод", level=1)
    document.add_paragraph(
        "Команде следует начать с операционно проверяемых шагов: FCR, время решения и отдельная работа с SMB. "
        "Каждый цикл следует завершать повторным измерением на том же наборе метрик."
    )
    document.save(path)


def write_executive_summary(path: Path, metrics: dict) -> None:
    document = Document()
    set_doc_defaults(document, "Резюме для руководства")
    document.add_heading("Удовлетворённость клиентов", 0)
    document.add_paragraph("Краткое резюме для руководства")
    document.add_heading("Вывод для руководства", level=1)
    document.add_paragraph(
        f"По {metrics['survey_count']} ответам общая удовлетворённость равна {metrics['overall_satisfaction']:.2f}/5, а NPS — {metrics['nps']:.0f}. "
        f"В {metrics['support_count']} обращениях поддержки FCR составил {metrics['support_fcr']:.0%}, среднее время решения — {metrics['average_resolution']:.2f} часа."
    )
    document.add_heading("Главный разрыв", level=1)
    document.add_paragraph(
        "SMB отстаёт от других сегментов: это самая низкая удовлетворённость 4.00 и NPS -30. "
        "Для сравнения: Enterprise — 4.50 и NPS 100; Mid-Market — 4.25 и NPS 85."
    )
    document.add_heading("Три приоритета", level=1)
    add_bullet(document, "Повысить FCR с 75% до 85% за 90 дней.")
    add_bullet(document, "Сократить среднее время решения с 3.40 до 2.00 часа за 120 дней.")
    add_bullet(document, "Запустить программу SMB и повысить удовлетворённость с 4.00 до 4.20 за 180 дней.")
    document.add_heading("Ограничение данных", level=1)
    document.add_paragraph(
        "В опросе нет Customer_ID, поэтому 50 ответов нельзя построчно связать с 20 обращениями. "
        "Резюме не делает выводов о корреляции на уровне клиента или о причинной связи; результаты нужно повторно измерить на новой волне."
    )
    document.save(path)


def solve(input_dir: Path, output_dir: Path) -> None:
    metrics = calculate(input_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    write_analysis(output_dir / ANALYSIS_NAME, metrics)
    actions = write_actions(output_dir / ACTION_NAME, metrics)
    write_report(output_dir / REPORT_NAME, metrics, actions)
    write_executive_summary(output_dir / SUMMARY_NAME, metrics)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=Path("/app/data/input_files"))
    parser.add_argument("--output-dir", type=Path, default=Path("/app/output"))
    args = parser.parse_args()
    solve(args.input_dir, args.output_dir)


if __name__ == "__main__":
    main()
