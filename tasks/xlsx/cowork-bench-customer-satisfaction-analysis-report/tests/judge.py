#!/usr/bin/env python3
"""从原始问卷和客服 CSV 重算真值的确定性 Harbor Judge。"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


ANALYSIS_NAME = "satisfaction_analysis.xlsx"
ACTION_NAME = "action_plans.xlsx"
REPORT_NAME = "satisfaction_report.docx"
SUMMARY_NAME = "executive_summary.docx"
OUTPUT_NAMES = [ANALYSIS_NAME, ACTION_NAME, REPORT_NAME, SUMMARY_NAME]
TASK_ID = "cowork-evo/cowork-bench-customer-satisfaction-analysis-report"
RUBRIC_VERSION = "cowork-bench-customer-satisfaction-analysis-report-v0.1"
JUDGE_VERSION = "cowork-bench-customer-satisfaction-analysis-report-v0.1"
PASS_THRESHOLD = 0.8
HURDLE_CAPS = {
    "R001": 0.0,
    "R002": 0.4,
    "R003": 0.55,
    "R004": 0.45,
    "R005": 0.7,
    "R006": 0.65,
    "R008": 0.5,
    "R009": 0.55,
    "R010": 0.6,
}
SEGMENT_ORDER = ["Enterprise", "Mid-Market", "SMB"]
SUMMARY_METRICS = [
    "Overall Satisfaction Score",
    "Net Promoter Score (NPS)",
    "First Contact Resolution Rate",
    "Average Resolution Time (Hours)",
]
DIMENSIONS = ["Product Quality", "Customer Service", "Value for Money", "Ease of Use", "Average Recommend Score"]
ACTION_HEADERS = ["Initiative", "Owner", "Target Metric", "Baseline", "Goal", "Timeline", "Status"]

CRITERION_SPECS = {
    "R001": ("四个指定交付物均可打开且工作表名正确", 5, "hurdle"),
    "R002": ("Summary Metrics 四项核心值与输入重算一致", 8, "hurdle"),
    "R003": ("Dimension Ratings 五个平均值正确", 6, "hurdle"),
    "R004": ("Segment Analysis 三个分群的样本数与指标正确", 8, "hurdle"),
    "R005": ("Action Plans 表头和至少三条完整行动正确", 4, "hurdle"),
    "R006": ("行动计划覆盖 FCR、解决时长和 SMB，基线与目标方向正确", 6, "hurdle"),
    "R007": ("详细报告为俄文可编辑文档且结构完整", 4, "positive"),
    "R008": ("详细报告的样本量、核心值、维度和分群与输入一致", 7, "hurdle"),
    "R009": ("高管摘要简洁且包含样本量、核心指标、SMB 结论和三项行动", 6, "hurdle"),
    "R010": ("报告声明无法逐客户关联，且没有上游无来源数字与声明", 5, "hurdle"),
    "R011": ("Excel 格式和 Word 标题结构达到专业交付要求", 4, "positive"),
    "R012": ("四个交付物可由 LibreOffice 实际打开并渲染", 3, "positive"),
    "R013": ("任一来源派生数字与输入重算真值冲突", -14, "penalty"),
    "R014": ("报告出现上游无来源样本量、相关系数、行业基准、预算或收入影响", -12, "penalty"),
    "R015": ("交付物缺失、损坏、截图化或核心数字跨文件冲突", -12, "penalty"),
}


def criterion(identifier: str, passed: bool, evidence: str = "") -> dict:
    description, weight, kind = CRITERION_SPECS[identifier]
    return {
        "criterion_id": identifier,
        "description": description,
        "weight": weight,
        "type": kind,
        "passed": bool(passed),
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
    }


def d(value) -> Decimal:
    return Decimal(str(value))


def round_2(value: Decimal) -> Decimal:
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def average(values) -> Decimal:
    items = [d(value) for value in values]
    return round_2(sum(items, Decimal("0")) / len(items))


def nps(rows: list[dict]) -> Decimal:
    promoters = sum(d(row["Recommend_Score"]) >= 9 for row in rows)
    detractors = sum(d(row["Recommend_Score"]) <= 6 for row in rows)
    return round_2(Decimal(promoters - detractors) * 100 / len(rows))


def derive_expected(input_dir: Path) -> dict:
    workbook = load_workbook(input_dir / "survey_template.xlsx", data_only=True)
    values = list(workbook["Responses"].iter_rows(values_only=True))
    headers = [str(value) for value in values[0]]
    responses = [dict(zip(headers, row)) for row in values[1:] if row and row[0]]
    workbook.close()
    with (input_dir / "support_data_export.csv").open(encoding="utf-8", newline="") as handle:
        tickets = list(csv.DictReader(handle))

    dimensions = {
        "Product Quality": average(row["Product_Quality"] for row in responses),
        "Customer Service": average(row["Customer_Service"] for row in responses),
        "Value for Money": average(row["Value_for_Money"] for row in responses),
        "Ease of Use": average(row["Ease_of_Use"] for row in responses),
        "Average Recommend Score": average(row["Recommend_Score"] for row in responses),
    }
    grouped = defaultdict(list)
    for row in responses:
        grouped[str(row["Customer_Type"])].append(row)
    segments = {}
    for segment in SEGMENT_ORDER:
        rows = grouped[segment]
        segments[segment] = {
            "Respondents": len(rows),
            "Overall_Satisfaction": average(row["Overall_Satisfaction"] for row in rows),
            "NPS": nps(rows),
            "First_Contact_Resolution_Rate": round_2(Decimal(sum(str(row["Resolved_First_Contact"]).lower() == "yes" for row in rows)) / len(rows)),
        }
    return {
        "survey_count": len(responses),
        "support_count": len(tickets),
        "summary": {
            "Overall Satisfaction Score": average(row["Overall_Satisfaction"] for row in responses),
            "Net Promoter Score (NPS)": nps(responses),
            "First Contact Resolution Rate": round_2(Decimal(sum(row["First_Contact_Resolution"].lower() == "yes" for row in tickets)) / len(tickets)),
            "Average Resolution Time (Hours)": average(row["Resolution_Time_Hours"] for row in tickets),
        },
        "dimensions": dimensions,
        "segments": segments,
    }


def normalized(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def close(actual, expected, tolerance: float = 0.011) -> bool:
    try:
        return abs(float(actual) - float(expected)) <= tolerance
    except (TypeError, ValueError):
        return False


def parse_number(value, percentage: bool = False) -> float | None:
    if isinstance(value, (int, float)):
        number = float(value)
        if percentage and number > 1.5:
            number /= 100
        return number
    text = str(value or "").strip().replace(",", ".")
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None
    number = float(match.group())
    if percentage and ("%" in text or number > 1.5):
        number /= 100
    return number


def find_header(sheet, headers: list[str]) -> int | None:
    wanted = [normalized(value) for value in headers]
    for row_number in range(1, sheet.max_row + 1):
        actual = [normalized(sheet.cell(row_number, column).value) for column in range(1, len(headers) + 1)]
        if actual == wanted:
            return row_number
    return None


def read_table(sheet, headers: list[str]) -> tuple[list[dict], int | None]:
    header_row = find_header(sheet, headers)
    if header_row is None:
        return [], None
    rows = []
    for row_number in range(header_row + 1, sheet.max_row + 1):
        values = [sheet.cell(row_number, column).value for column in range(1, len(headers) + 1)]
        if not values[0]:
            break
        rows.append(dict(zip(headers, values)))
    return rows, header_row


def docx_payload(path: Path) -> tuple[str, list[str], int]:
    document = Document(path)
    pieces = []
    headings = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            pieces.append(text)
            if str(paragraph.style.name).lower().startswith("heading") or paragraph.style.name == "Title":
                headings.append(text)
    for table in document.tables:
        for row in table.rows:
            pieces.extend(cell.text.strip() for cell in row.cells if cell.text.strip())
    with zipfile.ZipFile(path) as archive:
        media = sum(1 for name in archive.namelist() if name.startswith("word/media/") and not name.endswith("/"))
    return "\n".join(pieces), headings, media


def cyrillic_ratio(text: str) -> float:
    letters = [character for character in text if character.isalpha()]
    if not letters:
        return 0.0
    return sum("\u0400" <= character <= "\u04ff" for character in letters) / len(letters)


def labeled_value(text: str, label_pattern: str, value_pattern: str, distance: int = 80) -> bool:
    return re.search(label_pattern + rf".{{0,{distance}}}" + value_pattern, text, re.IGNORECASE | re.DOTALL) is not None


def factual_report_checks(text: str, expected: dict, detailed: bool) -> tuple[bool, list[str]]:
    compact = text.replace("\u00a0", " ")
    errors = []
    sample_ok = (
        re.search(r"\b50\b.{0,40}(?:ответ|опрос)", compact, re.IGNORECASE | re.DOTALL)
        and re.search(r"\b20\b.{0,45}(?:обращ|поддерж)", compact, re.IGNORECASE | re.DOTALL)
    )
    if not sample_ok:
        errors.append("缺真实样本量 50/20")
    required_labeled = [
        (r"(?:overall satisfaction|общ\w* удовлетвор)", r"4[.,]2(?:0)?"),
        (r"\bNPS\b", r"\b42\b"),
        (r"(?:first contact|\bFCR\b|перв\w* контакт)", r"75\s*%"),
        (r"(?:average resolution|средн\w* врем\w* решен)", r"3[.,]4(?:0)?"),
    ]
    for label, value in required_labeled:
        if not labeled_value(compact, label, value, 110):
            errors.append(f"缺指标 {label} -> {value}")
    smb_ok = labeled_value(compact, r"\bSMB\b", r"4[.,]0(?:0)?", 150) and labeled_value(compact, r"\bSMB\b", r"-30\b", 220)
    if not smb_ok:
        errors.append("缺 SMB=4.00/NPS -30")
    if detailed:
        dimension_values = {
            "Product Quality": "4.30",
            "Customer Service": "4.00",
            "Value for Money": "3.90",
            "Ease of Use": "4.40",
            "Average Recommend Score": "8.34",
        }
        for label, value in dimension_values.items():
            pattern = re.escape(value).replace("\\.", "[.,]")
            if not labeled_value(compact, re.escape(label), pattern, 70):
                errors.append(f"缺维度 {label}={value}")
        segment_values = {"Enterprise": ("10", "4.50", "100"), "Mid-Market": ("20", "4.25", "85"), "SMB": ("20", "4.00", "-30")}
        for segment, (count, satisfaction, segment_nps) in segment_values.items():
            pattern = re.escape(satisfaction).replace("\\.", "[.,]")
            if not labeled_value(compact, re.escape(segment), rf"\b{count}\b", 70):
                errors.append(f"缺分群样本 {segment}={count}")
            if not labeled_value(compact, re.escape(segment), pattern, 130):
                errors.append(f"缺分群满意度 {segment}={satisfaction}")
            if not labeled_value(compact, re.escape(segment), rf"NPS\s*{re.escape(segment_nps)}\b", 180):
                errors.append(f"缺分群 NPS {segment}={segment_nps}")
    return not errors, errors


def unsupported_claims(text: str) -> list[str]:
    checks = {
        "245 份问卷": r"\b245\b",
        "847 次互动": r"\b847\b",
        "r=0.68": r"\br\s*=\s*0[.,]68\b",
        "行业均值": r"industry average|средн\w* по отрасл",
        "150000 预算": r"150[\s,$]*000",
        "1.2-1.5 million": r"1[.,]2\s*[-–]\s*1[.,]5\s*(?:million|млн)",
        "8-12 percentage points": r"8\s*[-–]\s*12\s*(?:percentage|процент)",
    }
    return [label for label, pattern in checks.items() if re.search(pattern, text, re.IGNORECASE)]


def libreoffice_render(paths: list[Path], root: Path) -> tuple[dict[str, str], str]:
    libreoffice = shutil.which("libreoffice")
    pdftotext = shutil.which("pdftotext")
    if libreoffice is None or pdftotext is None:
        return {}, "缺少 libreoffice 或 pdftotext"
    source_dir = root / "source"
    pdf_dir = root / "pdf"
    source_dir.mkdir()
    pdf_dir.mkdir()
    profile_uri = (root / "lo-profile").resolve().as_uri()
    texts = {}
    for source in paths:
        local = source_dir / source.name
        shutil.copy2(source, local)
        converted = subprocess.run(
            [libreoffice, "--headless", f"-env:UserInstallation={profile_uri}", "--convert-to", "pdf", "--outdir", str(pdf_dir), str(local)],
            capture_output=True,
            text=True,
            timeout=300,
            check=False,
        )
        pdf_path = pdf_dir / f"{source.stem}.pdf"
        if converted.returncode != 0 or not pdf_path.is_file():
            return texts, (converted.stderr or converted.stdout or f"{source.name} 渲染失败")[-1000:]
        extracted = subprocess.run([pdftotext, "-raw", str(pdf_path), "-"], capture_output=True, text=True, timeout=120, check=False)
        if extracted.returncode != 0:
            return texts, (extracted.stderr or f"{source.name} PDF 文本失败")[-1000:]
        texts[source.name] = extracted.stdout
    return texts, "四个交付物的 LibreOffice PDF 渲染与文本抽取成功"


def evaluate(artifact_dir: Path) -> dict:
    paths = {name: artifact_dir / name for name in OUTPUT_NAMES}
    results = []
    if not all(path.is_file() for path in paths.values()):
        results.append(criterion("R001", False, f"文件存在性={{{', '.join(f'{name}: {path.is_file()}' for name, path in paths.items())}}}"))
        return build_result(results, list(paths.values()))

    input_dir = Path(__file__).resolve().parents[1] / "data" / "input_files"
    expected = derive_expected(input_dir)
    try:
        for path in paths.values():
            if path.suffix in (".xlsx", ".docx") and not zipfile.is_zipfile(path):
                raise ValueError(f"{path.name} 不是有效 OOXML ZIP")
        analysis_book = load_workbook(paths[ANALYSIS_NAME], data_only=True)
        action_book = load_workbook(paths[ACTION_NAME], data_only=True)
        analysis_formula_book = load_workbook(paths[ANALYSIS_NAME], data_only=False)
        action_formula_book = load_workbook(paths[ACTION_NAME], data_only=False)
        analysis_sheets = analysis_book.sheetnames
        action_sheets = action_book.sheetnames
        if "Satisfaction Analysis" not in analysis_sheets or "Action Plans" not in action_sheets:
            raise ValueError(f"工作表错误：{analysis_sheets} / {action_sheets}")
        analysis_sheet = analysis_book["Satisfaction Analysis"]
        action_sheet = action_book["Action Plans"]
        summary_rows, summary_header = read_table(analysis_sheet, ["Metric", "Value", "Source"])
        dimension_rows, dimension_header = read_table(analysis_sheet, ["Dimension", "Average", "Scale"])
        segment_rows, segment_header = read_table(
            analysis_sheet,
            ["Customer_Type", "Respondents", "Overall_Satisfaction", "NPS", "First_Contact_Resolution_Rate"],
        )
        action_rows, action_header = read_table(action_sheet, ACTION_HEADERS)
        analysis_format = analysis_formula_book["Satisfaction Analysis"]
        action_format = action_formula_book["Action Plans"]
        formatting = {
            "analysis_tables": len(analysis_format.tables),
            "action_tables": len(action_format.tables),
            "analysis_freeze": str(analysis_format.freeze_panes or ""),
            "action_freeze": str(action_format.freeze_panes or ""),
            "analysis_header_bold": all(analysis_format.cell(summary_header, column).font.bold for column in range(1, 4)) if summary_header else False,
            "action_header_bold": all(action_format.cell(action_header, column).font.bold for column in range(1, 8)) if action_header else False,
            "analysis_widths": all((analysis_format.column_dimensions[column].width or 0) >= 15 for column in "ABCDE"),
            "action_widths": all((action_format.column_dimensions[column].width or 0) >= 15 for column in "ABCDEFG"),
        }
        analysis_book.close()
        action_book.close()
        analysis_formula_book.close()
        action_formula_book.close()
        report_text, report_headings, report_media = docx_payload(paths[REPORT_NAME])
        executive_text, executive_headings, executive_media = docx_payload(paths[SUMMARY_NAME])
    except Exception as error:
        results.append(criterion("R001", False, f"无法读取交付物：{error}"))
        return build_result(results, list(paths.values()))

    structure_ok = (
        analysis_sheets == ["Satisfaction Analysis"]
        and action_sheets == ["Action Plans"]
        and summary_header is not None
        and dimension_header is not None
        and segment_header is not None
        and action_header is not None
        and len(summary_rows) == 4
        and len(dimension_rows) == 5
        and len(segment_rows) == 3
        and len(action_rows) >= 3
        and bool(report_text.strip())
        and bool(executive_text.strip())
    )
    results.append(
        criterion(
            "R001",
            structure_ok,
            f"工作表={analysis_sheets}/{action_sheets}，表头行={summary_header}/{dimension_header}/{segment_header}/{action_header}，行数={len(summary_rows)}/{len(dimension_rows)}/{len(segment_rows)}/{len(action_rows)}",
        )
    )

    actual_summary = {str(row["Metric"] or "").strip(): row for row in summary_rows}
    summary_errors = []
    for metric in SUMMARY_METRICS:
        actual = actual_summary.get(metric)
        if actual is None:
            summary_errors.append((metric, "缺失", float(expected["summary"][metric])))
            continue
        tolerance = 0.001 if metric == "First Contact Resolution Rate" else 0.011
        if not close(actual["Value"], expected["summary"][metric], tolerance):
            summary_errors.append((metric, actual["Value"], float(expected["summary"][metric])))
        source = normalized(actual["Source"])
        if metric in ("Overall Satisfaction Score", "Net Promoter Score (NPS)") and not ("survey" in source and "50" in source):
            summary_errors.append((metric, "Source 未指向 survey n=50", actual["Source"]))
        if metric in ("First Contact Resolution Rate", "Average Resolution Time (Hours)") and not ("support" in source and "20" in source):
            summary_errors.append((metric, "Source 未指向 support n=20", actual["Source"]))
    summary_ok = list(actual_summary) == SUMMARY_METRICS and not summary_errors
    results.append(criterion("R002", summary_ok, f"指标={list(actual_summary)}，错误={summary_errors}"))

    actual_dimensions = {str(row["Dimension"] or "").strip(): row for row in dimension_rows}
    dimension_errors = []
    for dimension in DIMENSIONS:
        actual = actual_dimensions.get(dimension)
        if actual is None or not close(actual["Average"] if actual else None, expected["dimensions"][dimension]):
            dimension_errors.append((dimension, actual["Average"] if actual else "缺失", float(expected["dimensions"][dimension])))
        expected_scale = "0-10" if dimension == "Average Recommend Score" else "1-5"
        if actual is not None and normalized(actual["Scale"]) != expected_scale:
            dimension_errors.append((dimension, "Scale", actual["Scale"], expected_scale))
    dimensions_ok = list(actual_dimensions) == DIMENSIONS and not dimension_errors
    results.append(criterion("R003", dimensions_ok, f"维度={list(actual_dimensions)}，错误={dimension_errors}"))

    actual_segments = {str(row["Customer_Type"] or "").strip(): row for row in segment_rows}
    segment_errors = []
    for segment in SEGMENT_ORDER:
        actual = actual_segments.get(segment)
        expected_row = expected["segments"][segment]
        if actual is None:
            segment_errors.append((segment, "缺失"))
            continue
        for field in ("Respondents", "Overall_Satisfaction", "NPS", "First_Contact_Resolution_Rate"):
            tolerance = 0 if field == "Respondents" else 0.011
            if not close(actual[field], expected_row[field], tolerance):
                segment_errors.append((segment, field, actual[field], float(expected_row[field])))
    segments_ok = list(actual_segments) == SEGMENT_ORDER and not segment_errors
    results.append(criterion("R004", segments_ok, f"分群={list(actual_segments)}，错误={segment_errors}"))

    action_completeness = len(action_rows) >= 3 and all(
        all(row.get(header) is not None and str(row.get(header)).strip() != "" for header in ACTION_HEADERS)
        for row in action_rows
    )
    results.append(criterion("R005", action_completeness, f"行数={len(action_rows)}，完整={action_completeness}"))

    family_rows = {}
    for row in action_rows:
        text = normalized(f"{row['Initiative']} {row['Target Metric']}")
        if "first contact" in text or re.search(r"\bfcr\b", text):
            family_rows.setdefault("fcr", row)
        elif "smb" in text and ("satisfaction" in text or "удовлетв" in text):
            family_rows.setdefault("smb", row)
        elif "resolution" in text and ("time" in text or "hours" in text):
            family_rows.setdefault("resolution", row)
    action_errors = []
    fcr_row = family_rows.get("fcr")
    if not fcr_row or not close(parse_number(fcr_row["Baseline"], True), expected["summary"]["First Contact Resolution Rate"], 0.001) or not (parse_number(fcr_row["Goal"], True) or 0) > float(expected["summary"]["First Contact Resolution Rate"]):
        action_errors.append(("fcr", fcr_row))
    resolution_row = family_rows.get("resolution")
    if not resolution_row or not close(parse_number(resolution_row["Baseline"]), expected["summary"]["Average Resolution Time (Hours)"], 0.011) or not (parse_number(resolution_row["Goal"]) is not None and parse_number(resolution_row["Goal"]) < float(expected["summary"]["Average Resolution Time (Hours)"])):
        action_errors.append(("resolution", resolution_row))
    smb_row = family_rows.get("smb")
    if not smb_row or not close(parse_number(smb_row["Baseline"]), expected["segments"]["SMB"]["Overall_Satisfaction"], 0.011) or not (parse_number(smb_row["Goal"]) or 0) > float(expected["segments"]["SMB"]["Overall_Satisfaction"]):
        action_errors.append(("smb", smb_row))
    role_keywords = ("head", "lead", "manager", "director", "team", "department", "руковод", "отдел", "команд")
    owner_errors = [row["Owner"] for row in action_rows if not any(keyword in normalized(row["Owner"]) for keyword in role_keywords)]
    invented_names = [name for name in ("Michael Chen", "Lisa Anderson") if any(normalized(row["Owner"]) == normalized(name) for row in action_rows)]
    actions_ok = not action_errors and not owner_errors and not invented_names
    results.append(criterion("R006", actions_ok, f"家族={list(family_rows)}，基线/目标错误={action_errors}，非角色Owner={owner_errors}，上游虚构人名={invented_names}"))

    report_norm = normalized(report_text)
    report_section_tokens = ("резюм", "метод", "метрик", "измерен", "сегмент", "улучш", "огранич")
    report_structure_ok = len(report_text) >= 1800 and cyrillic_ratio(report_text) >= 0.55 and len(report_headings) >= 7 and all(token in report_norm for token in report_section_tokens) and report_media == 0
    results.append(criterion("R007", report_structure_ok, f"字符={len(report_text)}，西里尔比={cyrillic_ratio(report_text):.2f}，标题={len(report_headings)}，媒体={report_media}"))

    report_factual_ok, report_factual_errors = factual_report_checks(report_text, expected, detailed=True)
    results.append(criterion("R008", report_factual_ok, f"错误={report_factual_errors}"))

    executive_factual_ok, executive_factual_errors = factual_report_checks(executive_text, expected, detailed=False)
    executive_norm = normalized(executive_text)
    action_tokens = [
        ("fcr" in executive_norm or "first contact" in executive_norm or "первого контакт" in executive_norm),
        ("resolution" in executive_norm or "время решен" in executive_norm),
        ("smb" in executive_norm and ("програм" in executive_norm or "program" in executive_norm)),
    ]
    smb_low = "smb" in executive_norm and any(token in executive_norm for token in ("отста", "самая низк", "lowest", "underperform"))
    executive_ok = (
        executive_factual_ok
        and 650 <= len(executive_text) <= 2500
        and cyrillic_ratio(executive_text) >= 0.5
        and len(executive_headings) >= 4
        and all(action_tokens)
        and smb_low
        and executive_media == 0
    )
    results.append(criterion("R009", executive_ok, f"字符={len(executive_text)}，错误={executive_factual_errors}，行动={action_tokens}，SMB最低={smb_low}，媒体={executive_media}"))

    combined_reports = report_text + "\n" + executive_text
    unsupported = unsupported_claims(combined_reports)
    limitation_checks = []
    for text in (report_text, executive_text):
        low = normalized(text)
        limitation_checks.append(
            "customer_id" in low
            and ("нельзя" in low or "невозмож" in low)
            and "корреляц" in low
            and "причин" in low
        )
    limitations_ok = all(limitation_checks) and not unsupported
    results.append(criterion("R010", limitations_ok, f"限制声明={limitation_checks}，无来源声明={unsupported}"))

    formatting_ok = (
        formatting["analysis_tables"] >= 3
        and formatting["action_tables"] >= 1
        and formatting["analysis_freeze"] == "A3"
        and formatting["action_freeze"] == "A2"
        and formatting["analysis_header_bold"]
        and formatting["action_header_bold"]
        and formatting["analysis_widths"]
        and formatting["action_widths"]
        and len(report_headings) >= 7
        and len(executive_headings) >= 4
    )
    results.append(criterion("R011", formatting_ok, json.dumps(formatting, ensure_ascii=False)))

    with tempfile.TemporaryDirectory(prefix="cowork-evo-csat-report-") as temp_dir:
        rendered_texts, render_evidence = libreoffice_render(list(paths.values()), Path(temp_dir))
    render_checks = {
        ANALYSIS_NAME: all(token in normalized(rendered_texts.get(ANALYSIS_NAME, "")) for token in ("overall satisfaction", "nps", "smb")),
        ACTION_NAME: all(token in normalized(rendered_texts.get(ACTION_NAME, "")) for token in ("initiative", "owner", "first contact", "smb")),
        REPORT_NAME: all(token in normalized(rendered_texts.get(REPORT_NAME, "")) for token in ("удовлетвор", "nps", "smb")),
        SUMMARY_NAME: all(token in normalized(rendered_texts.get(SUMMARY_NAME, "")) for token in ("руковод", "nps", "smb")),
    }
    render_ok = len(rendered_texts) == 4 and all(render_checks.values())
    results.append(criterion("R012", render_ok, f"{render_evidence}，检查={render_checks}"))

    source_error = not (summary_ok and dimensions_ok and segments_ok and report_factual_ok and executive_factual_ok)
    unsupported_error = bool(unsupported)
    cross_conflict = not (structure_ok and render_ok and summary_ok and report_factual_ok and executive_factual_ok)
    results.append(criterion("R013", source_error, f"摘要={summary_ok}，维度={dimensions_ok}，分群={segments_ok}，详报={report_factual_ok}，高管={executive_factual_ok}"))
    results.append(criterion("R014", unsupported_error, f"无来源声明={unsupported}"))
    results.append(criterion("R015", cross_conflict, f"结构={structure_ok}，渲染={render_ok}，Excel核心={summary_ok}，详报={report_factual_ok}，高管={executive_factual_ok}"))
    return build_result(results, list(paths.values()))


def build_result(results: list[dict], artifact_paths: list[Path]) -> dict:
    present = {item["criterion_id"] for item in results}
    for identifier in CRITERION_SPECS:
        if identifier not in present:
            results.append(criterion(identifier, False, "因交付不存在、损坏或门槛失败，本判据未执行"))
    results.sort(key=lambda item: item["criterion_id"])
    positive_total = sum(item["weight"] for item in results if item["weight"] > 0)
    earned = sum(item["weight"] for item in results if item["weight"] > 0 and item["passed"])
    penalties = sum(item["weight"] for item in results if item["weight"] < 0 and item["passed"])
    reward = max(0.0, min(1.0, (earned + penalties) / positive_total))
    hurdle_failures = [item["criterion_id"] for item in results if item["type"] == "hurdle" and not item["passed"]]
    if hurdle_failures:
        reward = min(reward, min(HURDLE_CAPS[identifier] for identifier in hurdle_failures))
    digests = [hashlib.sha256(path.read_bytes()).hexdigest() for path in artifact_paths if path.is_file()]
    if not digests:
        digests = [hashlib.sha256(b"").hexdigest()]
    criterion_results = []
    for item in results:
        score = (0.0 if item["passed"] else 1.0) if item["type"] == "penalty" else item["score"]
        criterion_results.append(
            {
                "criterion_id": item["criterion_id"],
                "score": score,
                "evidence": item["evidence"],
                "method": item["method"],
                "raw": {"description": item["description"], "weight": item["weight"], "type": item["type"], "passed": item["passed"]},
            }
        )
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": digests,
        "criterion_results": criterion_results,
        "reward": round(reward, 6),
        "passed": reward >= PASS_THRESHOLD,
        "hurdle_failures": hurdle_failures,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path)
    parser.add_argument("--reward-file", type=Path)
    args = parser.parse_args()
    result = evaluate(args.artifact_dir)
    rendered = json.dumps(result, ensure_ascii=False, indent=2)
    print(rendered)
    if args.result:
        args.result.parent.mkdir(parents=True, exist_ok=True)
        args.result.write_text(rendered + "\n", encoding="utf-8")
    if args.reward_file:
        args.reward_file.parent.mkdir(parents=True, exist_ok=True)
        args.reward_file.write_text(f"{result['reward']:.6f}\n", encoding="utf-8")
    raise SystemExit(0 if result["passed"] else 1)


if __name__ == "__main__":
    main()
