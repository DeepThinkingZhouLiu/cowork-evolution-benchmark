#!/usr/bin/env python3
"""Independent deterministic Judge for the offline faculty workload review."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import subprocess
import zipfile
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation


TASK_ID = "cowork-evo/cowork-bench-faculty-workload-review"
RUBRIC_VERSION = "cowork-bench-faculty-workload-review-v0.1"
JUDGE_VERSION = "cowork-bench-faculty-workload-review-deterministic-v0.1"
OUTPUTS = ["Faculty_Workload.xlsx", "Faculty_Workload_Analysis.xlsx", "Workload_Review.pptx"]
EXPECTED_INPUT_HASHES = {
    "faculty_directory.csv": "c799754e494d9edbc33992852f1aec9c771b1e15dd545b20d1cd64e4f1f19476",
    "workload_policy.pdf": "8def78ba5d60a69ce6fa88590562dc030bf98cdf0b1d470d85a478baad52ccbc",
    "workload_standards.json": "a02b95499e815bcfad690616613f9ccf45db3163af7292e3cd74894cbddf9127",
    "canvas_courses.csv": "3563f2cf7d6b0ab84c6188a7ec42ccee2ca80817fc80ae0a0cbeef4b47a47c44",
    "canvas_teacher_enrollments.csv": "4ad16a382fcbef47ee7dfeb638319cc95e92c2a3784c10ec59db1c5b12964b20",
    "canvas_student_enrollments.csv": "88219cc8cbb84f2954be106616fabc64717c1ac5330bc332dea7f3c989607682",
    "canvas_assignments.csv": "6c34feb35e7bdc49173e9cd9de56dae6411960409ceafa4d0227437e5f1c4d60",
    "canvas_snapshot_manifest.json": "79aae9b667522d92081be64b0894c472bec1055e7a1f83a0a751ba66b70de96c",
}
INSTRUCTOR_HEADERS = ["Instructor", "Courses_Count", "Total_Students", "Total_Assignments", "Est_Grading_Hours", "Weekly_Grading_Hours", "Overloaded_YN"]
DEPARTMENT_HEADERS = ["Department", "Instructor_Count", "Course_Count", "Total_Students"]
WEIGHTS = {"R001": 0.12, "R002": 0.22, "R003": 0.16, "R004": 0.12, "R005": 0.14, "R006": 0.10, "R007": 0.08, "R008": 0.06}


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def one_decimal(value: Decimal) -> float:
    return float(value.quantize(Decimal("0.1"), rounding=ROUND_HALF_UP))


def criterion(cid: str, score: bool | float, evidence: str) -> dict:
    return {"criterion_id": cid, "score": float(score), "evidence": evidence, "method": "deterministic"}


def derive_truth(input_dir: Path):
    with (input_dir / "workload_standards.json").open(encoding="utf-8") as stream:
        standards = json.load(stream)
    with (input_dir / "canvas_courses.csv").open(encoding="utf-8", newline="") as stream:
        courses = {row["Course_Code"]: row for row in csv.DictReader(stream)}
    with (input_dir / "canvas_teacher_enrollments.csv").open(encoding="utf-8", newline="") as stream:
        teacher_rows = list(csv.DictReader(stream))
    student_counts = defaultdict(int)
    with (input_dir / "canvas_student_enrollments.csv").open(encoding="utf-8", newline="") as stream:
        for row in csv.DictReader(stream):
            student_counts[row["Course_ID"]] += 1
    assignment_counts = defaultdict(int)
    with (input_dir / "canvas_assignments.csv").open(encoding="utf-8", newline="") as stream:
        for row in csv.DictReader(stream):
            assignment_counts[row["Course_ID"]] += 1

    by_instructor = defaultdict(list)
    for row in teacher_rows:
        by_instructor[row["Instructor"]].append(row)
    instructors = []
    for instructor in sorted(by_instructor, key=lambda name: min(row["Login_ID"] for row in by_instructor[name])):
        rows = by_instructor[instructor]
        course_ids = {row["Course_ID"] for row in rows}
        students = sum(student_counts[row["Course_ID"]] for row in rows)
        assignments = sum(assignment_counts[row["Course_ID"]] for row in rows)
        grading = Decimal(students) * Decimal(str(standards["grading_hours_per_student"]))
        weekly = grading / Decimal(str(standards["semester_weeks"]))
        overloaded = len(course_ids) > int(standards["max_courses_per_semester"]) or weekly > Decimal(str(standards["max_weekly_hours"]))
        instructors.append([instructor, len(course_ids), students, assignments, one_decimal(grading), one_decimal(weekly), "Yes" if overloaded else "No"])

    by_department = defaultdict(lambda: {"instructors": set(), "courses": set(), "students": 0})
    for row in teacher_rows:
        course = courses[row["Course_Code"]]
        bucket = by_department[course["Subject_Area"]]
        bucket["instructors"].add(row["Instructor"])
        bucket["courses"].add(row["Course_Code"])
        bucket["students"] += student_counts[row["Course_ID"]]
    departments = [[department, len(bucket["instructors"]), len(bucket["courses"]), bucket["students"]] for department, bucket in sorted(by_department.items())]
    overloaded = sum(row[-1] == "Yes" for row in instructors)
    return instructors, departments, overloaded


def rows_from_sheet(sheet):
    return [[cell.value for cell in row] for row in sheet.iter_rows()]


def values_match(actual, expected) -> bool:
    if len(actual) != len(expected):
        return False
    for actual_row, expected_row in zip(actual, expected):
        if len(actual_row) != len(expected_row):
            return False
        for a, e in zip(actual_row, expected_row):
            if isinstance(e, float):
                try:
                    if abs(float(a) - e) > 0.11:
                        return False
                except (TypeError, ValueError):
                    return False
            elif a != e:
                return False
    return True


def presentation_text(path: Path) -> tuple[int, str]:
    prs = Presentation(path)
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                parts.append(shape.text)
            if getattr(shape, "has_table", False):
                parts.extend(cell.text for row in shape.table.rows for cell in row.cells)
    text = "\n".join(parts)
    return len(prs.slides), text


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    input_dir = Path(__file__).parents[1] / "data/input_files"
    results = []
    try:
        instructors, departments, overloaded = derive_truth(input_dir)
        expected_instructor = [INSTRUCTOR_HEADERS] + instructors
        expected_department = [DEPARTMENT_HEADERS] + departments
        actual_files = sorted(path.name for path in args.artifact_dir.iterdir() if path.is_file()) if args.artifact_dir.is_dir() else []
        missing = sorted(set(OUTPUTS) - set(actual_files))
        extras = sorted(set(actual_files) - set(OUTPUTS))
        native = all((args.artifact_dir / name).is_file() and zipfile.is_zipfile(args.artifact_dir / name) for name in OUTPUTS)
        results.append(criterion("R001", native and not missing and not extras, f"missing={missing}; extras={extras}; OOXML={native}"))

        workbook = load_workbook(args.artifact_dir / "Faculty_Workload.xlsx", data_only=True)
        load_ok = "Instructor Load" in workbook.sheetnames and "Department Summary" in workbook.sheetnames
        output_instructor = rows_from_sheet(workbook["Instructor Load"]) if load_ok else []
        output_department = rows_from_sheet(workbook["Department Summary"]) if load_ok else []
        results.append(criterion("R002", load_ok and values_match(output_instructor, expected_instructor), f"rows={len(output_instructor)-1}; expected={len(instructors)}"))
        results.append(criterion("R003", load_ok and values_match(output_department, expected_department), f"departments={output_department[1:] if output_department else []}"))

        analysis = load_workbook(args.artifact_dir / "Faculty_Workload_Analysis.xlsx", data_only=True)
        analysis_rows = rows_from_sheet(analysis["Instructor Load"]) if "Instructor Load" in analysis.sheetnames else []
        analysis_ok = analysis.properties.title == "Faculty Workload Analysis" and values_match(analysis_rows, expected_instructor)
        results.append(criterion("R004", analysis_ok, f"sheets={analysis.sheetnames}; rows={len(analysis_rows)-1}"))

        slide_count, slide_text = presentation_text(args.artifact_dir / "Workload_Review.pptx")
        lower = slide_text.lower()
        dept_text_ok = all(row[0].lower() in lower and str(row[1]) in slide_text and str(row[3]) in slide_text for row in departments)
        ppt_ok = slide_count >= 4 and "faculty workload" in lower and re.search(r"total instructors\s*:\s*33", lower) and re.search(r"overloaded instructors\s*:\s*23", lower) and "redistribution" in lower and dept_text_ok
        results.append(criterion("R005", bool(ppt_ok), f"slides={slide_count}; departments={dept_text_ok}; headline counts={len(instructors)}/{overloaded}"))

        cross_ok = analysis_ok and bool(ppt_ok) and all(str(value) in slide_text for row in departments for value in row[1:])
        results.append(criterion("R006", cross_ok, "workbook and presentation share independently derived counts and department totals"))

        policy_text = subprocess.run(["pdftotext", str(input_dir / "workload_policy.pdf"), "-"], check=True, capture_output=True, text=True).stdout.lower()
        input_hashes_ok = all(digest(input_dir / name) == expected for name, expected in EXPECTED_INPUT_HASHES.items())
        rules_ok = all(token in policy_text for token in ("4", "40", "0,5", "16-недель")) and int(json.loads((input_dir / "workload_standards.json").read_text(encoding="utf-8"))["max_courses_per_semester"]) == 4
        results.append(criterion("R007", input_hashes_ok and rules_ok, f"input_hashes={input_hashes_ok}; policy_rules={rules_ok}"))

        professional = load_ok and all(bool(workbook[name].tables) and workbook[name].freeze_panes == "A2" for name in ("Instructor Load", "Department Summary")) and all(workbook["Instructor Load"].cell(row, 5).number_format == "0.0" for row in range(2, workbook["Instructor Load"].max_row + 1))
        professional = professional and slide_count >= 4 and len(slide_text) > 500
        results.append(criterion("R008", professional, "editable Excel tables, frozen headers, numeric formats, and text-based presentation"))
        workbook.close()
        analysis.close()
    except Exception as exc:
        results = [criterion("R001", False, f"Judge read failure: {exc}")]

    by_id = {item["criterion_id"]: item["score"] for item in results}
    reward = sum(WEIGHTS[cid] * by_id.get(cid, 0.0) for cid in WEIGHTS)
    passed = reward >= 0.75 and all(by_id.get(cid) == 1.0 for cid in ("R001", "R002", "R003", "R004", "R005"))
    artifact_hashes = [digest(args.artifact_dir / name) for name in OUTPUTS if (args.artifact_dir / name).is_file()]
    payload = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": artifact_hashes, "criterion_results": results, "reward": round(reward, 6), "passed": passed}
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.write_text(f"{reward:.6f}\n", encoding="utf-8")
    raise SystemExit(0 if passed else 1)


if __name__ == "__main__":
    main()
