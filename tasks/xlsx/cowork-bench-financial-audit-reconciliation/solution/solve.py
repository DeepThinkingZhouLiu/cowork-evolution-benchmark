#!/usr/bin/env python3
"""Cowork-Bench 财务审计对账的逐题离线参考解法。"""

from __future__ import annotations

import argparse
import csv
import hashlib
import subprocess
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

EXPECTED_HASHES = {
    "authoritative_ledger.csv": "6d141be371f2b2d8942e493b1f59b4a631ccd0eb482bb2ced6bcc3dd4989cb00",
    "q1_manual_tracking.xlsx": "d980f32983abca1921f64282f9fb6d347e2080dcc7ac9ee721f7a8c43719d95e",
    "system_mapping_guide.pdf": "457fd78134f2ea0877129436a13bb1cbce79270bf9f991d57a3139c03f61fddf",
}
VALID_STATUSES = {"completed", "pending"}


def digest(path: Path) -> str:
    value = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            value.update(block)
    return value.hexdigest()


def read_inputs(input_dir: Path) -> tuple[dict[str, dict], list[dict]]:
    for name, expected in EXPECTED_HASHES.items():
        if digest(input_dir / name) != expected:
            raise ValueError(f"输入哈希不匹配：{name}")
    guide = subprocess.run(["pdftotext", str(input_dir / "system_mapping_guide.pdf"), "-"], check=True, capture_output=True, text=True).stdout.lower()
    if not all(token in guide for token in ("order id", "1 день", "amount", "completed", "pending")):
        raise ValueError("映射指南缺少匹配规则证据")
    with (input_dir / "authoritative_ledger.csv").open(encoding="utf-8-sig", newline="") as stream:
        ledger = {row["Order ID"]: row for row in csv.DictReader(stream)}
    workbook = load_workbook(input_dir / "q1_manual_tracking.xlsx", data_only=True)
    sheet = workbook["Manual Tracking"]
    headers = [cell.value for cell in sheet[1]]
    manual = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    workbook.close()
    return ledger, manual


def reconcile(ledger: dict[str, dict], manual: list[dict]) -> list[dict]:
    records = []
    for row in manual:
        order_id = str(row["Order ID"])
        source = ledger.get(order_id)
        reasons = []
        if source is None:
            reasons.append("missing from authoritative ledger")
        else:
            manual_date = date.fromisoformat(str(row["Date"])[:10])
            ledger_date = date.fromisoformat(str(source["Date"])[:10])
            if abs((manual_date - ledger_date).days) > 1:
                reasons.append("date mismatch")
            if Decimal(str(row["Amount"])).quantize(Decimal("0.01")) != Decimal(str(source["Amount"])).quantize(Decimal("0.01")):
                reasons.append("amount mismatch")
            if str(row["Status"]).strip().lower() not in VALID_STATUSES or str(source["Status"]).strip().lower() not in VALID_STATUSES:
                reasons.append("invalid status")
        records.append({"order_id": order_id, "status": "discrepancy" if reasons else "matched", "reasons": reasons, "manual": row, "ledger": source})
    return records


def style_sheet(sheet, table_name: str, ref: str) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = ref
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)


def write_excel(records: list[dict], total_amount: Decimal, output_dir: Path) -> None:
    workbook = Workbook()
    reconciliation = workbook.active
    reconciliation.title = "Reconciliation"
    reconciliation.append(["Order ID", "Match Status"])
    for record in records:
        reconciliation.append([record["order_id"], record["status"]])
    style_sheet(reconciliation, "ReconciliationTable", f"A1:B{reconciliation.max_row}")
    reconciliation.column_dimensions["A"].width = 16
    reconciliation.column_dimensions["B"].width = 20

    summary = workbook.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    matched = sum(record["status"] == "matched" for record in records)
    discrepancies = len(records) - matched
    for row in (("Total Orders", len(records)), ("Matched", matched), ("Discrepancies", discrepancies), ("Total Amount", float(total_amount))):
        summary.append(row)
    summary["B5"].number_format = "#,##0.00"
    style_sheet(summary, "AuditSummary", f"A1:B{summary.max_row}")
    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 16
    workbook.save(output_dir / "Audit_Reconciliation_Report.xlsx")
    workbook.close()


def discrepancy_detail(record: dict) -> str:
    source = record["ledger"]
    manual = record["manual"]
    if source is None:
        return "missing from authoritative ledger"
    if "date mismatch" in record["reasons"]:
        return f"date mismatch: manual {manual['Date']}, ledger {source['Date']}"
    if "amount mismatch" in record["reasons"]:
        return f"amount mismatch: manual {Decimal(str(manual['Amount'])):.2f}, ledger {Decimal(str(source['Amount'])):.2f}"
    return "; ".join(record["reasons"])


def write_word(records: list[dict], total_amount: Decimal, output_dir: Path) -> None:
    matched = sum(record["status"] == "matched" for record in records)
    discrepancies = [record for record in records if record["status"] == "discrepancy"]
    document = Document()
    document.add_heading("Financial Audit Reconciliation Report", 0)
    document.add_heading("Executive Summary", level=1)
    document.add_paragraph(f"The audit reconciled {len(records)} manual orders against the authoritative ledger. Matched: {matched}. Discrepancies: {len(discrepancies)}. Total Amount: {total_amount:.2f} RUB.")
    document.add_heading("Discrepancies", level=1)
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.rows[0].cells[0].text = "Order ID"
    table.rows[0].cells[1].text = "Reason"
    for record in discrepancies:
        cells = table.add_row().cells
        cells[0].text = record["order_id"]
        cells[1].text = discrepancy_detail(record)
    document.add_heading("Method", level=1)
    document.add_paragraph("Orders were matched by exact Order ID, date difference of no more than one day, exact amount, and case-insensitive completed/pending status validation.")
    document.save(output_dir / "Audit_Report.docx")


def write_email(records: list[dict], output_dir: Path) -> None:
    matched = sum(record["status"] == "matched" for record in records)
    discrepancy_ids = [record["order_id"] for record in records if record["status"] == "discrepancy"]
    lines = [
        "From: audit@company.invalid",
        "To: finance@company.com",
        "Subject: Financial Audit Reconciliation - Review Required",
        "Date: Sat, 29 Aug 2026 10:00:00 +0800",
        "MIME-Version: 1.0",
        "Content-Type: text/plain; charset=utf-8",
        "X-CoworkEvoBench-State: offline-draft",
        "",
        "Offline draft - not sent.",
        f"Matched: {matched}. Discrepancies: {len(discrepancy_ids)}.",
        "Discrepancy orders: " + ", ".join(discrepancy_ids) + ".",
        "Please review the attached report: Audit_Report.docx.",
    ]
    (output_dir / "audit_notification.eml").write_text("\r\n".join(lines) + "\r\n", encoding="utf-8", newline="")


def write_calendar(records: list[dict], output_dir: Path) -> None:
    matched = sum(record["status"] == "matched" for record in records)
    discrepancy_ids = [record["order_id"] for record in records if record["status"] == "discrepancy"]
    description = f"Offline draft. Matched: {matched}; Discrepancies: {len(discrepancy_ids)}; Orders: {', '.join(discrepancy_ids)}"
    lines = [
        "BEGIN:VCALENDAR",
        "VERSION:2.0",
        "PRODID:-//CoworkEvoBench//Financial Audit//EN",
        "CALSCALE:GREGORIAN",
        "METHOD:PUBLISH",
        "BEGIN:VEVENT",
        "UID:financial-audit-reconciliation-review@coworkevobench.local",
        "DTSTAMP:20260829T020000Z",
        "DTSTART:20260903T100000",
        "DTEND:20260903T110000",
        "SUMMARY:Financial Audit Reconciliation Review",
        "DESCRIPTION:" + description.replace(",", "\\,"),
        "END:VEVENT",
        "END:VCALENDAR",
        "",
    ]
    (output_dir / "audit_review_meeting.ics").write_text("\r\n".join(lines), encoding="utf-8", newline="")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    ledger, manual = read_inputs(args.input_dir)
    records = reconcile(ledger, manual)
    total_amount = sum((Decimal(str(row["Amount"])) for row in manual), Decimal("0")).quantize(Decimal("0.01"))
    write_excel(records, total_amount, args.output_dir)
    write_word(records, total_amount, args.output_dir)
    write_email(records, args.output_dir)
    write_calendar(records, args.output_dir)


if __name__ == "__main__":
    main()
