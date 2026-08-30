#!/usr/bin/env python3
"""确定性 Judge：从台账、手工表和 PDF 规则独立重算并检查四个交付物。"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import subprocess
import zipfile
from datetime import date, datetime
from decimal import Decimal
from email import policy
from email.parser import BytesParser
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

TASK_ID = "cowork-evo/cowork-bench-financial-audit-reconciliation"
RUBRIC_VERSION = "cowork-bench-financial-audit-reconciliation-v0.1"
JUDGE_VERSION = "cowork-bench-financial-audit-reconciliation-deterministic-v0.1"
OUTPUTS = ["Audit_Reconciliation_Report.xlsx", "Audit_Report.docx", "audit_notification.eml", "audit_review_meeting.ics"]
EXPECTED_HASHES = {"authoritative_ledger.csv": "6d141be371f2b2d8942e493b1f59b4a631ccd0eb482bb2ced6bcc3dd4989cb00", "q1_manual_tracking.xlsx": "d980f32983abca1921f64282f9fb6d347e2080dcc7ac9ee721f7a8c43719d95e", "system_mapping_guide.pdf": "457fd78134f2ea0877129436a13bb1cbce79270bf9f991d57a3139c03f61fddf"}
VALID_STATUSES = {"completed", "pending"}
WEIGHTS = {"R001": 0.10, "R002": 0.20, "R003": 0.15, "R004": 0.10, "R005": 0.15, "R006": 0.10, "R007": 0.05, "R008": 0.05, "R009": 0.05, "R010": 0.05}


def digest(path: Path) -> str:
    value = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            value.update(block)
    return value.hexdigest()


def criterion(cid: str, score: bool | float, evidence: str) -> dict:
    return {"criterion_id": cid, "score": float(score), "evidence": evidence, "method": "deterministic"}


def derive_truth(input_dir: Path) -> tuple[list[dict], Decimal, bool]:
    guide = subprocess.run(["pdftotext", str(input_dir / "system_mapping_guide.pdf"), "-"], check=True, capture_output=True, text=True).stdout.lower()
    guide_ok = all(token in guide for token in ("order id", "1 день", "amount", "completed", "pending"))
    with (input_dir / "authoritative_ledger.csv").open(encoding="utf-8-sig", newline="") as stream:
        ledger = {row["Order ID"]: row for row in csv.DictReader(stream)}
    workbook = load_workbook(input_dir / "q1_manual_tracking.xlsx", data_only=True)
    sheet = workbook["Manual Tracking"]
    headers = [cell.value for cell in sheet[1]]
    manual = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    workbook.close()
    records = []
    for row in manual:
        order_id = str(row["Order ID"])
        source = ledger.get(order_id)
        reasons = []
        if source is None:
            reasons.append("missing from authoritative ledger")
        else:
            if abs((date.fromisoformat(str(row["Date"])[:10]) - date.fromisoformat(str(source["Date"])[:10])).days) > 1:
                reasons.append("date mismatch")
            if Decimal(str(row["Amount"])).quantize(Decimal("0.01")) != Decimal(str(source["Amount"])).quantize(Decimal("0.01")):
                reasons.append("amount mismatch")
            if str(row["Status"]).lower() not in VALID_STATUSES or str(source["Status"]).lower() not in VALID_STATUSES:
                reasons.append("invalid status")
        records.append({"order_id": order_id, "status": "discrepancy" if reasons else "matched", "reasons": reasons})
    total = sum((Decimal(str(row["Amount"])) for row in manual), Decimal("0")).quantize(Decimal("0.01"))
    return records, total, guide_ok


def docx_content(path: Path) -> tuple[str, list[str], list[list[list[str]]]]:
    document = Document(path)
    text = [paragraph.text for paragraph in document.paragraphs]
    headings = [paragraph.text.strip().lower() for paragraph in document.paragraphs if paragraph.style.name.startswith("Heading") or paragraph.style.name == "Title"]
    tables = []
    for table in document.tables:
        table_rows = []
        for row in table.rows:
            values = [cell.text for cell in row.cells]
            text.extend(values)
            table_rows.append(values)
        tables.append(table_rows)
    return "\n".join(text), headings, tables


def parse_ics(path: Path) -> tuple[list[dict[str, str]], bool]:
    text = path.read_text(encoding="utf-8")
    calendar_ok = "BEGIN:VCALENDAR" in text and "END:VCALENDAR" in text
    events = []
    for body in re.findall(r"BEGIN:VEVENT\r?\n(.*?)\r?\nEND:VEVENT", text, flags=re.S):
        props = {}
        for line in body.splitlines():
            if ":" in line:
                key, value = line.split(":", 1)
                props[key.split(";", 1)[0].upper()] = value.replace("\\,", ",")
        events.append(props)
    return events, calendar_ok


def numeric_mentions(text: str) -> bool:
    compact = re.sub(r"[^0-9]", "", text)
    return "326398" in compact and re.search(r"(?:match|匹配|сверен)[^0-9]{0,20}11", text, re.I) is not None and re.search(r"(?:discrep|差异|расхожд)[^0-9]{0,20}4", text, re.I) is not None


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    input_dir = Path(__file__).parents[1] / "data/input_files"
    results: list[dict] = []
    try:
        missing = [name for name in OUTPUTS if not (args.artifact_dir / name).is_file()]
        native = not missing and zipfile.is_zipfile(args.artifact_dir / OUTPUTS[0]) and zipfile.is_zipfile(args.artifact_dir / OUTPUTS[1])
        records, total_amount, guide_ok = derive_truth(input_dir)
        discrepancy_reasons = {record["order_id"]: record["reasons"][0] for record in records if record["reasons"]}
        discrepancy_ids = list(discrepancy_reasons)
        results.append(criterion("R001", native, f"缺失={missing}; Office OOXML={native}"))

        workbook = load_workbook(args.artifact_dir / "Audit_Reconciliation_Report.xlsx", data_only=True)
        if not {"Reconciliation", "Summary"}.issubset(workbook.sheetnames):
            raise ValueError("审计工作簿缺少指定工作表")
        reconciliation = workbook["Reconciliation"]
        recon_headers = [str(cell.value).strip() for cell in reconciliation[1]]
        order_index = recon_headers.index("Order ID") + 1
        status_index = recon_headers.index("Match Status") + 1
        output_recon = [(str(reconciliation.cell(row, order_index).value).strip(), str(reconciliation.cell(row, status_index).value).strip().lower()) for row in range(2, reconciliation.max_row + 1) if reconciliation.cell(row, order_index).value is not None]
        expected_recon = [(record["order_id"], record["status"]) for record in records]
        results.append(criterion("R002", output_recon == expected_recon, f"期望={expected_recon}; 输出={output_recon}"))

        summary = workbook["Summary"]
        summary_values = {str(summary.cell(row, 1).value).strip(): summary.cell(row, 2).value for row in range(2, summary.max_row + 1) if summary.cell(row, 1).value is not None}
        summary_ok = float(summary_values.get("Total Orders", -1)) == 15 and float(summary_values.get("Matched", -1)) == 11 and float(summary_values.get("Discrepancies", -1)) == 4 and abs(float(summary_values.get("Total Amount", -1)) - float(total_amount)) <= 0.01
        results.append(criterion("R003", summary_ok, f"Summary={summary_values}"))

        report_text, headings, tables = docx_content(args.artifact_dir / "Audit_Report.docx")
        sections_ok = any("executive summary" in heading or "резюме" in heading for heading in headings) and any("discrep" in heading or "расхожд" in heading for heading in headings)
        results.append(criterion("R004", sections_ok and numeric_mentions(report_text), f"标题={headings}; 摘要数字检查"))
        report_reasons = {}
        for table in tables:
            for row in table[1:]:
                if len(row) >= 2 and re.fullmatch(r"WC-[0-9]{3}", row[0].strip()):
                    report_reasons[row[0].strip()] = row[1].strip().lower()
        reason_ok = set(report_reasons) == set(discrepancy_reasons) and all(reason in report_reasons[order_id] for order_id, reason in discrepancy_reasons.items())
        results.append(criterion("R005", reason_ok, f"差异原因={discrepancy_reasons}"))

        message = BytesParser(policy=policy.default).parsebytes((args.artifact_dir / "audit_notification.eml").read_bytes())
        email_body = message.get_body(preferencelist=("plain",)).get_content() if message.is_multipart() else message.get_content()
        email_text = str(email_body)
        email_ok = "finance@company.com" in str(message.get("To", "")).lower() and "Financial Audit Reconciliation" in str(message.get("Subject", "")) and ("draft" in email_text.lower() or "未发送" in email_text) and "Audit_Report.docx" in email_text and all(order_id in email_text for order_id in discrepancy_ids) and re.search(r"Matched[^0-9]{0,10}11", email_text, re.I) and re.search(r"Discrepancies[^0-9]{0,10}4", email_text, re.I)
        results.append(criterion("R006", bool(email_ok), "邮件收件人、主题、草稿标记、摘要和报告提及检查"))

        events, calendar_ok = parse_ics(args.artifact_dir / "audit_review_meeting.ics")
        event = events[0] if len(events) == 1 else {}
        required_event = {"UID", "DTSTART", "DTEND", "SUMMARY", "DESCRIPTION"}.issubset(event)
        chronology = False
        if required_event:
            chronology = datetime.strptime(event["DTEND"].rstrip("Z"), "%Y%m%dT%H%M%S") > datetime.strptime(event["DTSTART"].rstrip("Z"), "%Y%m%dT%H%M%S")
        event_text = event.get("SUMMARY", "") + " " + event.get("DESCRIPTION", "")
        event_ok = calendar_ok and required_event and chronology and "Financial Audit Reconciliation Review" in event.get("SUMMARY", "") and all(order_id in event_text for order_id in discrepancy_ids) and re.search(r"Matched[^0-9]{0,10}11", event_text, re.I) and re.search(r"Discrepancies[^0-9]{0,10}4", event_text, re.I)
        results.append(criterion("R007", bool(event_ok), f"VEVENT={len(events)}; 时间顺序={chronology}"))

        cross_ok = output_recon == expected_recon and summary_ok and all(order_id in report_text and order_id in email_text and order_id in event_text for order_id in discrepancy_ids) and numeric_mentions(report_text)
        results.append(criterion("R008", cross_ok, "四个交付物的 11/4 和四个差异订单一致"))
        professional = bool(reconciliation.tables) and bool(summary.tables) and reconciliation.freeze_panes == "A2" and summary.freeze_panes == "A2" and len(headings) >= 3 and len(tables) >= 1
        results.append(criterion("R009", professional, f"Excel 表格={bool(reconciliation.tables) and bool(summary.tables)}; Word 标题={len(headings)}; 表格={len(tables)}"))
        hashes_ok = all(digest(input_dir / name) == expected for name, expected in EXPECTED_HASHES.items())
        results.append(criterion("R010", hashes_ok and guide_ok, f"输入哈希={hashes_ok}; PDF 规则={guide_ok}"))
        workbook.close()
    except Exception as exc:
        results = [criterion("R001", False, f"Judge 读取失败：{exc}")]
    by_id = {item["criterion_id"]: item["score"] for item in results}
    reward = sum(WEIGHTS[cid] * by_id.get(cid, 0.0) for cid in WEIGHTS)
    passed = reward >= 0.8 and all(by_id.get(cid) == 1.0 for cid in ("R001", "R002", "R003", "R004", "R005", "R006", "R007", "R008"))
    artifacts = [digest(args.artifact_dir / name) for name in OUTPUTS if (args.artifact_dir / name).is_file()]
    payload = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": artifacts, "criterion_results": results, "reward": reward, "passed": passed}
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.write_text(f"{reward:.6f}\n", encoding="utf-8")
    raise SystemExit(0 if passed else 1)


if __name__ == "__main__":
    main()
