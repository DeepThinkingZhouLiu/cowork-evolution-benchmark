#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import hashlib
import subprocess
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

CONTRACT = {"authoritative_ledger.csv": "6d141be371f2b2d8942e493b1f59b4a631ccd0eb482bb2ced6bcc3dd4989cb00", "q1_manual_tracking.xlsx": "d980f32983abca1921f64282f9fb6d347e2080dcc7ac9ee721f7a8c43719d95e", "system_mapping_guide.pdf": "457fd78134f2ea0877129436a13bb1cbce79270bf9f991d57a3139c03f61fddf"}
VALID = {"completed", "pending"}


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_facts(input_dir: Path):
    for name, expected in CONTRACT.items():
        if digest(input_dir / name) != expected: raise ValueError(f"input hash mismatch: {name}")
    with (input_dir / "authoritative_ledger.csv").open(encoding="utf-8-sig", newline="") as stream:
        ledger = {row["Order ID"]: row for row in csv.DictReader(stream)}
    source = load_workbook(input_dir / "q1_manual_tracking.xlsx", data_only=True)
    sheet = source["Manual Tracking"]
    headers = [cell.value for cell in sheet[1]]
    manual = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    source.close()
    records = []
    for row in manual:
        order = str(row["Order ID"]); other = ledger.get(order); reasons = []
        if other is None: reasons.append("missing from authoritative ledger")
        else:
            if abs((date.fromisoformat(str(row["Date"])[:10]) - date.fromisoformat(str(other["Date"])[:10])).days) > 1: reasons.append("date mismatch")
            if Decimal(str(row["Amount"])).quantize(Decimal("0.01")) != Decimal(str(other["Amount"])).quantize(Decimal("0.01")): reasons.append("amount mismatch")
            if str(row["Status"]).strip().lower() not in VALID or str(other["Status"]).strip().lower() not in VALID: reasons.append("invalid status")
        records.append({"order": order, "status": "discrepancy" if reasons else "matched", "reason": "; ".join(reasons), "manual": row, "ledger": other})
    return ledger, manual, records


def style(ws, table_name: str, cols: int):
    ws.freeze_panes = "A2"; ws.sheet_view.showGridLines = False
    for cell in ws[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F4E78"); cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=cols):
        for cell in row: cell.alignment = Alignment(vertical="top", wrap_text=True)
    if ws.max_row >= 2:
        tab = Table(displayName=table_name, ref=f"A1:{chr(64 + cols)}{ws.max_row}"); tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False); ws.add_table(tab)
    for column in ws.columns:
        letter = column[0].column_letter; ws.column_dimensions[letter].width = min(42, max(14, max(len(str(c.value or "")) for c in column) + 2))


def main() -> None:
    parser = argparse.ArgumentParser(); parser.add_argument("--input-dir", type=Path, required=True); parser.add_argument("--output-dir", type=Path, required=True); args = parser.parse_args()
    ledger, manual, records = load_facts(args.input_dir)
    wb = Workbook(); source = wb.active; source.title = "Source Ledger"
    source.append(["Order ID", "Date", "Amount", "Status"])
    for row in ledger.values(): source.append([row["Order ID"], row["Date"], float(row["Amount"]), row["Status"]])
    style(source, "SourceLedger", 4)
    tracking = wb.create_sheet("Manual Tracking")
    tracking.append(["Order ID", "Date", "Amount", "Customer", "Status"])
    for row in manual: tracking.append([row["Order ID"], row["Date"], float(row["Amount"]), row["Customer"], row["Status"]])
    style(tracking, "ManualTracking", 5)
    reconciliation = wb.create_sheet("Reconciliation"); reconciliation.append(["Order ID", "Match Status", "Reason"])
    for rec in records: reconciliation.append([rec["order"], rec["status"], rec["reason"]])
    style(reconciliation, "Reconciliation", 3)
    summary = wb.create_sheet("Summary"); summary.append(["Metric", "Value"])
    summary.append(["Total Orders", "=COUNTA('Manual Tracking'!A2:A16)"])
    summary.append(["Matched", '=COUNTIF(Reconciliation!B2:B16,"matched")'])
    summary.append(["Discrepancies", '=COUNTIF(Reconciliation!B2:B16,"discrepancy")'])
    summary.append(["Total Amount", "=SUM('Manual Tracking'!C2:C16)"])
    style(summary, "Summary", 2); summary["B5"].number_format = "#,##0.00"
    details = wb.create_sheet("Discrepancy Details"); details.append(["Order ID", "Reason", "Manual Date", "Ledger Date", "Manual Amount", "Ledger Amount"])
    for rec in records:
        if rec["status"] != "discrepancy": continue
        other = rec["ledger"]; manual_row = rec["manual"]
        details.append([rec["order"], rec["reason"], manual_row["Date"], other["Date"] if other else "", float(manual_row["Amount"]), float(other["Amount"]) if other else ""])
    style(details, "DiscrepancyDetails", 6)
    matched = sum(r["status"] == "matched" for r in records); discrep = len(records) - matched; ids = ", ".join(r["order"] for r in records if r["status"] == "discrepancy")
    email = wb.create_sheet("Audit Email Draft"); email.append(["To", "Subject", "Body", "Status"]); email.append(["finance@company.com", "Financial Audit Reconciliation - Review Required", f"Offline draft - not sent. Matched: {matched}. Discrepancies: {discrep}. Discrepancy orders: {ids}.", "Draft - Not Sent"]); style(email, "AuditEmailDraft", 4)
    meeting = wb.create_sheet("Review Meeting Draft"); meeting.append(["Meeting Title", "Start", "End", "Description", "Status"]); start = date.today() + timedelta(days=7); end = start + timedelta(hours=1); meeting.append(["Financial Audit Reconciliation Review", start.strftime("%Y-%m-%d 10:00"), end.strftime("%Y-%m-%d 11:00"), f"Offline draft - not created. Matched: {matched}; Discrepancies: {discrep}; Orders: {ids}", "Draft - Not Created"]); style(meeting, "ReviewMeetingDraft", 5)
    wb.calculation.calcMode = "auto"; wb.calculation.fullCalcOnLoad = True; wb.calculation.forceFullCalc = True
    args.output_dir.mkdir(parents=True, exist_ok=True); wb.save(args.output_dir / "financial_audit_reconciliation.xlsx"); wb.close()


if __name__ == "__main__": main()
