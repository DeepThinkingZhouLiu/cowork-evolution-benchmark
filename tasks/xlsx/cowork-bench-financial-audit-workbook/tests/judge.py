#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

OUTPUT = "financial_audit_reconciliation.xlsx"
SHEETS = ["Source Ledger", "Manual Tracking", "Reconciliation", "Summary", "Discrepancy Details", "Audit Email Draft", "Review Meeting Draft"]
CONTRACT = {"authoritative_ledger.csv": "6d141be371f2b2d8942e493b1f59b4a631ccd0eb482bb2ced6bcc3dd4989cb00", "q1_manual_tracking.xlsx": "d980f32983abca1921f64282f9fb6d347e2080dcc7ac9ee721f7a8c43719d95e", "system_mapping_guide.pdf": "457fd78134f2ea0877129436a13bb1cbce79270bf9f991d57a3139c03f61fddf"}
VALID = {"completed", "pending"}
WEIGHTS = {"R001": 10, "R002": 15, "R003": 25, "R004": 15, "R005": 10, "R006": 10, "R007": 5, "R008": 5, "R009": 5, "R010": -20}
KINDS = {"R001": "hurdle", "R002": "hurdle", "R003": "hurdle", "R004": "hurdle", "R005": "hurdle", "R006": "hurdle", "R007": "positive", "R008": "positive", "R009": "positive", "R010": "penalty"}
CAPS = {"R001": 0.0, "R002": 0.5, "R003": 0.45, "R004": 0.6, "R005": 0.65, "R006": 0.7}


def digest(path: Path) -> str: return hashlib.sha256(path.read_bytes()).hexdigest()


def facts(source: Path):
    guide = subprocess.run(["pdftotext", str(source / "system_mapping_guide.pdf"), "-"], check=True, capture_output=True, text=True).stdout.lower()
    guide_ok = all(token in guide for token in ("order id", "1 день", "amount", "completed", "pending"))
    with (source / "authoritative_ledger.csv").open(encoding="utf-8-sig", newline="") as stream: ledger = list(csv.DictReader(stream))
    source_wb = load_workbook(source / "q1_manual_tracking.xlsx", data_only=True); ws = source_wb["Manual Tracking"]; headers = [c.value for c in ws[1]]; manual = [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]; source_wb.close()
    by_id = {row["Order ID"]: row for row in ledger}; records = []
    for row in manual:
        other = by_id.get(str(row["Order ID"])); reasons = []
        if other is None: reasons.append("missing from authoritative ledger")
        else:
            if abs((date.fromisoformat(str(row["Date"])[:10]) - date.fromisoformat(str(other["Date"])[:10])).days) > 1: reasons.append("date mismatch")
            if Decimal(str(row["Amount"])).quantize(Decimal("0.01")) != Decimal(str(other["Amount"])).quantize(Decimal("0.01")): reasons.append("amount mismatch")
            if str(row["Status"]).strip().lower() not in VALID or str(other["Status"]).strip().lower() not in VALID: reasons.append("invalid status")
        records.append({"order": str(row["Order ID"]), "status": "discrepancy" if reasons else "matched", "reason": "; ".join(reasons), "manual": row, "ledger": other})
    return ledger, manual, records, guide_ok


def criterion(cid, score, evidence): return {"criterion_id": cid, "score": float(score), "evidence": evidence, "method": "deterministic", "weight": WEIGHTS[cid], "type": KINDS[cid]}


def normalized_rows(values):
    return [["" if value is None else (float(value) if isinstance(value, (int, float)) else value) for value in row] for row in values]


def recalc(path: Path) -> Path | None:
    binary = shutil.which("libreoffice") or shutil.which("soffice")
    if not binary: return None
    root = Path(tempfile.mkdtemp(prefix="financial-audit-recalc-")); src = root / "src"; out = root / "out"; src.mkdir(); out.mkdir(); shutil.copy2(path, src / path.name)
    run = subprocess.run([binary, "--headless", f"-env:UserInstallation={(root/'profile').resolve().as_uri()}", "--convert-to", "xlsx", "--outdir", str(out), str(src / path.name)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=240, check=False)
    result = out / path.name; return result if run.returncode == 0 and result.is_file() else None


def evaluate(artifact_dir: Path, do_render: bool) -> dict:
    source = Path(__file__).parents[1] / "data/input_files"; output = artifact_dir / OUTPUT; names = sorted(p.name for p in artifact_dir.iterdir()) if artifact_dir.is_dir() else []
    ledger, manual, records, guide_ok = facts(source); results = []
    valid = output.is_file() and zipfile.is_zipfile(output); wb = None; cached = recalc(output) if valid else None; cached_wb = load_workbook(cached, data_only=True) if cached else None
    if valid:
        try: wb = load_workbook(output, data_only=False); valid = wb.sheetnames == SHEETS
        except Exception: valid = False
    results.append(criterion("R001", valid and names == [OUTPUT], f"valid={valid}; files={names}; sheets={wb.sheetnames if wb else None}"))
    if not valid:
        for cid in ("R002", "R003", "R004", "R005", "R006"): results.append(criterion(cid, 0, "invalid workbook"))
        results += [criterion("R007", 0, "invalid workbook"), criterion("R008", 0, "invalid workbook"), criterion("R009", 1 if all(digest(source/n)==h for n,h in CONTRACT.items()) and guide_ok else 0, "input checks"), criterion("R010", 0, "invalid output")]
        return finalize(results, output)
    def table(name): return list(wb[name].iter_rows(values_only=True))
    source_rows = table("Source Ledger"); expected_source = [["Order ID", "Date", "Amount", "Status"]] + [[r["Order ID"], r["Date"], float(r["Amount"]), r["Status"]] for r in ledger]
    manual_rows = table("Manual Tracking"); expected_manual = [["Order ID", "Date", "Amount", "Customer", "Status"]] + [[r["Order ID"], r["Date"], float(r["Amount"]), r["Customer"], r["Status"]] for r in manual]
    source_ok = normalized_rows(source_rows) == normalized_rows(expected_source); manual_ok = normalized_rows(manual_rows) == normalized_rows(expected_manual)
    results.append(criterion("R002", 1 if source_ok and manual_ok else 0, f"source_copy={source_ok}; manual_copy={manual_ok}"))
    recon = table("Reconciliation"); recon_expected = [["Order ID", "Match Status", "Reason"]] + [[r["order"], r["status"], r["reason"]] for r in records]; recon_ok = normalized_rows(recon) == normalized_rows(recon_expected)
    results.append(criterion("R003", 1 if recon_ok else 0, f"independent 15-order reconciliation={recon_ok}"))
    summary = table("Summary"); smap = {str(r[0]): r[1] for r in summary[1:] if r and r[0] is not None}; matched = sum(r["status"] == "matched" for r in records); discrep = len(records)-matched; total = float(sum((Decimal(str(r["Amount"])) for r in manual), Decimal("0")))
    summary_ok = set(("Total Orders", "Matched", "Discrepancies", "Total Amount")).issubset(smap)
    if cached_wb:
        cs = cached_wb["Summary"]; vals = [cs.cell(i,2).value for i in range(2,6)]; summary_ok &= vals[0] == len(records) and vals[1] == matched and vals[2] == discrep and abs(float(vals[3]) - total) < .01
    summary_ok &= all(isinstance(smap[k], str) and smap[k].startswith("=") for k in ("Total Orders", "Matched", "Discrepancies", "Total Amount"))
    results.append(criterion("R004", 1 if summary_ok else 0, f"summary={summary_ok}; expected=15/11/4/{total:.2f}"))
    details = table("Discrepancy Details"); detail_expected = [["Order ID", "Reason", "Manual Date", "Ledger Date", "Manual Amount", "Ledger Amount"]] + [[r["order"], r["reason"], r["manual"]["Date"], r["ledger"]["Date"] if r["ledger"] else "", float(r["manual"]["Amount"]), float(r["ledger"]["Amount"]) if r["ledger"] else ""] for r in records if r["status"] == "discrepancy"]
    details_ok = normalized_rows(details) == normalized_rows(detail_expected)
    results.append(criterion("R005", 1 if details_ok else 0, f"four discrepancy detail rows={details_ok}"))
    email = table("Audit Email Draft"); meeting = table("Review Meeting Draft"); ids = [r["order"] for r in records if r["status"] == "discrepancy"]; email_ok = len(email) == 2 and email[1][0] == "finance@company.com" and "Financial Audit Reconciliation" in str(email[1][1]) and all(x in str(email[1][2]) for x in ["Matched: 11", "Discrepancies: 4", *ids]) and email[1][3] == "Draft - Not Sent"; meeting_ok = len(meeting) == 2 and "Financial Audit Reconciliation Review" in str(meeting[1][0]) and all(x in str(meeting[1][3]) for x in ["Matched: 11", "Discrepancies: 4", *ids]) and meeting[1][4] == "Draft - Not Created"
    results.append(criterion("R006", 1 if email_ok and meeting_ok else 0, f"email={email_ok}; meeting={meeting_ok}"))
    formulas = all(isinstance(wb["Summary"].cell(r,2).value, str) and wb["Summary"].cell(r,2).value.startswith("=") for r in range(2,6)); tables_ok = all(bool(wb[n].tables) for n in SHEETS[:5]); structure_ok = wb.sheetnames == SHEETS and all(wb[n].freeze_panes == "A2" for n in SHEETS[:5]) and getattr(wb.calculation, "calcMode", None) == "auto" and bool(getattr(wb.calculation, "fullCalcOnLoad", False)) and bool(getattr(wb.calculation, "forceFullCalc", False))
    results.append(criterion("R007", 1 if formulas and tables_ok and structure_ok else 0, f"formulas={formulas}; tables={tables_ok}; structure={structure_ok}"))
    render_ok = True; render_ev = "render skipped"
    if do_render:
        binary = shutil.which("libreoffice") or shutil.which("soffice"); pdfinfo = shutil.which("pdfinfo"); root = Path(tempfile.mkdtemp(prefix="financial-audit-render-")); profile = root / "profile"; run = subprocess.run([binary, "--headless", f"-env:UserInstallation={profile.resolve().as_uri()}", "--convert-to", "pdf", "--outdir", str(root), str(output)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=240, check=False) if binary else None; pdf = root / "financial_audit_reconciliation.pdf"; info = subprocess.run([pdfinfo, str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, check=False).stdout if pdfinfo and pdf.is_file() else ""; match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE); render_ok = bool(run and run.returncode == 0 and pdf.is_file() and match and int(match.group(1)) >= 1); render_ev = f"PDF={pdf.stat().st_size if pdf.is_file() else 0}; {match.group(0) if match else 'no pages'}"
    results.append(criterion("R008", 1 if render_ok else 0, render_ev))
    hashes_ok = all(digest(source/n) == h for n,h in CONTRACT.items()) and guide_ok; results.append(criterion("R009", 1 if hashes_ok else 0, f"input_sha256_and_pdf_rules={hashes_ok}"))
    unchanged = digest(output) in {digest(source/n) for n in CONTRACT}; results.append(criterion("R010", 1 if unchanged else 0, "unchanged input penalty triggered" if unchanged else "output differs from inputs"))
    if cached_wb: cached_wb.close()
    wb.close(); return finalize(results, output)


def finalize(results, output):
    positive = sum(w for w in WEIGHTS.values() if w > 0); earned = sum(item["score"] * WEIGHTS[item["criterion_id"]] for item in results if WEIGHTS[item["criterion_id"]] > 0); penalties = sum(WEIGHTS[item["criterion_id"]] for item in results if WEIGHTS[item["criterion_id"]] < 0 and item["score"] >= 1); reward = max(0.0, min(1.0, (earned + penalties) / positive)); failures = [item["criterion_id"] for item in results if KINDS[item["criterion_id"]] == "hurdle" and item["score"] < 1];
    if failures: reward = min(reward, min(CAPS[c] for c in failures))
    return {"task_id":"cowork-evo/cowork-bench-financial-audit-workbook","rubric_version":"cowork-bench-financial-audit-workbook-v0.1","judge_version":"2.0.0","artifact_sha256":[digest(output)] if output.is_file() else [],"criterion_results":[{k:item[k] for k in ("criterion_id","score","evidence","method")} for item in sorted(results,key=lambda x:x["criterion_id"])],"reward":round(reward,6),"passed":reward >= .8 and not failures,"hurdle_failures":failures}


def main():
    p=argparse.ArgumentParser(); p.add_argument("--artifact-dir",type=Path,required=True); p.add_argument("--result",type=Path,required=True); p.add_argument("--reward-file",type=Path,required=True); p.add_argument("--skip-render",action="store_true"); a=p.parse_args(); result=evaluate(a.artifact_dir.resolve(),not a.skip_render); a.result.parent.mkdir(parents=True,exist_ok=True); a.result.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8"); a.reward_file.parent.mkdir(parents=True,exist_ok=True); a.reward_file.write_text(f"{result['reward']:.6f}\n"); print(json.dumps(result,ensure_ascii=False,indent=2)); raise SystemExit(0 if result["passed"] else 1)


if __name__ == "__main__": main()
