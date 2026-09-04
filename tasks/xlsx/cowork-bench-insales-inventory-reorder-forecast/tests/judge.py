#!/usr/bin/env python3
"""Cowork-Bench 库存补货预测题的确定性 Harbor Judge。"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import shutil
import subprocess
import tempfile
import zipfile
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook


OUTPUT_NAME = "Inventory_Forecast.xlsx"
TASK_ID = "cowork-evo/cowork-bench-insales-inventory-reorder-forecast"
RUBRIC_VERSION = "cowork-bench-insales-inventory-static-v0.1"
JUDGE_VERSION = "cowork-bench-insales-inventory-static-v0.1"
HURDLE_CAPS = {"R001": 0.0, "R002": 0.74, "R003": 0.74, "R006": 0.74, "R007": 0.74, "R008": 0.74}

STOCK_HEADERS = [
    "Product_ID", "Product_Name", "Category", "Current_Stock", "Avg_Daily_Sales",
    "Days_Until_Safety_Stock", "Lead_Time", "Needs_Reorder", "Reorder_Date",
]
SCHEDULE_HEADERS = ["Product_Name", "Category", "Current_Stock", "Reorder_Date", "Lead_Time_Days"]

CRITERION_SPECS = {
    "R001": ("指定 XLSX 可打开并包含两张规定工作表及表头", 3, "hurdle"),
    "R002": ("Stock Analysis 完整且准确覆盖 77 个来源产品", 5, "hurdle"),
    "R003": ("按 361 天窗口正确计算 77 个产品的平均日销量", 6, "hurdle"),
    "R004": ("77 行交付期均与供应商品类映射一致", 4, "positive"),
    "R005": ("安全库存可售天数在 77 行上均计算正确", 5, "positive"),
    "R006": ("Needs_Reorder 精确标记 18 个应补货产品", 5, "hurdle"),
    "R007": ("18 个补货日期均由固定快照日和交付期正确推导", 6, "hurdle"),
    "R008": ("Reorder Schedule 与 Stock Analysis 的补货集合及字段一一对应", 5, "hurdle"),
    "R009": ("工作簿可由 LibreOffice 真实打开、重算和重存", 3, "positive"),
    "R010": ("快照产品存在遗漏、伪造或关键原值错误", -10, "penalty"),
    "R011": ("补货集合、日期或两表对应关系存在错误", -10, "penalty"),
    "R012": ("交付损坏、非可编辑工作簿或无法由 LibreOffice 打开", -10, "penalty"),
}


def criterion(identifier: str, passed: bool, evidence: str = "") -> dict:
    description, weight, kind = CRITERION_SPECS[identifier]
    return {
        "criterion_id": identifier,
        "description": description,
        "weight": weight,
        "type": kind,
        "passed": bool(passed),
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
    }


def load_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def normalized(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def to_float(value) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def to_int(value) -> int | None:
    number = to_float(value)
    if number is None or not number.is_integer():
        return None
    return int(number)


def parse_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    match = re.search(r"(\d{4})-(\d{2})-(\d{2})", str(value or ""))
    if not match:
        return None
    try:
        return date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
    except ValueError:
        return None


def name_matches(actual, expected: str) -> bool:
    left = normalized(actual)
    right = normalized(expected)
    if not left or not right:
        return False
    if left == right:
        return True
    shorter = min(len(left), len(right))
    return shorter >= 40 and (left.startswith(right) or right.startswith(left))


def derive_expected(input_dir: Path) -> tuple[dict[int, dict], int, int]:
    products = load_json(input_dir / "store_products.json")
    orders = load_json(input_dir / "store_orders.json")
    lead_times = load_json(input_dir / "supplier_lead_times.json")
    snapshot_date = date.fromisoformat((input_dir / "snapshot_date.txt").read_text(encoding="utf-8").strip())

    eligible_orders = [order for order in orders if order["status"] in {"completed", "processing"}]
    order_dates = [datetime.fromisoformat(order["order_date"]).date() for order in eligible_orders]
    history_days = (max(order_dates) - min(order_dates)).days
    sold = defaultdict(int)
    for order in eligible_orders:
        for item in order["line_items"]:
            sold[int(item["product_id"])] += int(item["quantity"])

    expected = {}
    for product in products:
        if product["status"] != "publish" or int(product["current_stock"]) <= 0:
            continue
        product_id = int(product["product_id"])
        category = product["categories"][0]["name"]
        stock = int(product["current_stock"])
        average = round(sold[product_id] / history_days, 4)
        days_until = 999 if average == 0 else round(max(stock - 5, 0) / average, 1)
        needs_reorder = sold[product_id] > 0 and days_until <= 90
        if needs_reorder:
            computed = snapshot_date + timedelta(days=math.floor(days_until) - int(lead_times[category]))
            reorder_date = max(snapshot_date + timedelta(days=1), computed)
        else:
            reorder_date = None
        expected[product_id] = {
            "product_id": product_id,
            "name": product["product_name"],
            "category": category,
            "stock": stock,
            "average": average,
            "days_until": days_until,
            "lead_time": int(lead_times[category]),
            "needs_reorder": needs_reorder,
            "reorder_date": reorder_date,
        }
    return expected, history_days, len(eligible_orders)


def roundtrip_with_libreoffice(path: Path, work_root: Path) -> tuple[Path | None, str]:
    libreoffice = shutil.which("libreoffice")
    if libreoffice is None:
        return None, "缺少 libreoffice"
    source_dir = work_root / "source"
    output_dir = work_root / "roundtrip"
    source_dir.mkdir()
    output_dir.mkdir()
    source_copy = source_dir / OUTPUT_NAME
    shutil.copy2(path, source_copy)
    profile_uri = (work_root / "lo-profile").resolve().as_uri()
    process = subprocess.run(
        [
            libreoffice,
            "--headless",
            f"-env:UserInstallation={profile_uri}",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(output_dir),
            str(source_copy),
        ],
        capture_output=True,
        text=True,
        timeout=300,
        check=False,
    )
    converted = output_dir / OUTPUT_NAME
    if process.returncode != 0 or not converted.is_file():
        detail = process.stderr or process.stdout or "LibreOffice 重存失败"
        return None, detail[-1000:]
    try:
        workbook = load_workbook(converted, data_only=True, read_only=True)
        sheet_names = workbook.sheetnames
        workbook.close()
    except Exception as error:
        return None, f"重存后无法读取：{error}"
    return converted, f"重存成功，工作表={sheet_names}"


def sheet_records(workbook, sheet_name: str, required_headers: list[str]) -> tuple[list[dict], list[str]]:
    if sheet_name not in workbook.sheetnames:
        return [], required_headers
    sheet = workbook[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return [], required_headers
    header = [str(value).strip() if value is not None else "" for value in rows[0]]
    index = {name: position for position, name in enumerate(header)}
    missing = [name for name in required_headers if name not in index]
    if missing:
        return [], missing
    records = []
    for row in rows[1:]:
        if not row or all(value is None or str(value).strip() == "" for value in row):
            continue
        records.append({name: row[index[name]] if index[name] < len(row) else None for name in required_headers})
    return records, []


def evaluate(artifact_dir: Path) -> dict:
    output_path = artifact_dir / OUTPUT_NAME
    results = []
    if not output_path.is_file():
        results.append(criterion("R001", False, f"缺少 {OUTPUT_NAME}"))
        return build_result(results, output_path)

    input_dir = Path(__file__).resolve().parents[1] / "data" / "input_files"
    expected, history_days, eligible_order_count = derive_expected(input_dir)

    with tempfile.TemporaryDirectory(prefix="cowork-evo-inventory-") as temp_dir:
        roundtrip_path, roundtrip_evidence = roundtrip_with_libreoffice(output_path, Path(temp_dir))
        load_path = roundtrip_path or output_path
        try:
            if not zipfile.is_zipfile(load_path):
                raise ValueError("不是有效 OOXML ZIP")
            workbook = load_workbook(load_path, data_only=True)
            stock_records, stock_missing = sheet_records(workbook, "Stock Analysis", STOCK_HEADERS)
            schedule_records, schedule_missing = sheet_records(workbook, "Reorder Schedule", SCHEDULE_HEADERS)
            workbook.close()
        except Exception as error:
            results.append(criterion("R001", False, f"XLSX 无法读取：{error}"))
            return build_result(results, output_path)

    structure_ok = not stock_missing and not schedule_missing
    results.append(
        criterion(
            "R001",
            structure_ok,
            f"Stock Analysis 缺失表头={stock_missing}，Reorder Schedule 缺失表头={schedule_missing}，行数={len(stock_records)}/{len(schedule_records)}",
        )
    )

    actual_by_id = {}
    duplicate_ids = []
    invalid_id_rows = []
    for record in stock_records:
        product_id = to_int(record["Product_ID"])
        if product_id is None:
            invalid_id_rows.append(record["Product_ID"])
            continue
        if product_id in actual_by_id:
            duplicate_ids.append(product_id)
        actual_by_id[product_id] = record

    expected_ids = set(expected)
    actual_ids = set(actual_by_id)
    missing_ids = sorted(expected_ids - actual_ids)
    extra_ids = sorted(actual_ids - expected_ids)
    source_errors = []
    for product_id in sorted(expected_ids & actual_ids):
        actual = actual_by_id[product_id]
        target = expected[product_id]
        if not name_matches(actual["Product_Name"], target["name"]):
            source_errors.append((product_id, "name", actual["Product_Name"] ))
        if normalized(actual["Category"]) != normalized(target["category"]):
            source_errors.append((product_id, "category", actual["Category"]))
        if to_int(actual["Current_Stock"]) != target["stock"]:
            source_errors.append((product_id, "stock", actual["Current_Stock"]))
    source_ok = (
        len(stock_records) == 77
        and not missing_ids
        and not extra_ids
        and not duplicate_ids
        and not invalid_id_rows
        and not source_errors
    )
    results.append(
        criterion(
            "R002",
            source_ok,
            f"行数={len(stock_records)}/77，缺失ID={missing_ids[:6]}，多余ID={extra_ids[:6]}，重复ID={duplicate_ids[:6]}，原值错误={source_errors[:6]}",
        )
    )

    average_errors = []
    lead_errors = []
    days_errors = []
    needs_errors = []
    date_errors = []
    actual_reorder_ids = set()
    for product_id in sorted(expected_ids & actual_ids):
        actual = actual_by_id[product_id]
        target = expected[product_id]
        average = to_float(actual["Avg_Daily_Sales"])
        if average is None or abs(average - target["average"]) > 0.00011:
            average_errors.append((product_id, actual["Avg_Daily_Sales"], target["average"]))
        if to_int(actual["Lead_Time"]) != target["lead_time"]:
            lead_errors.append((product_id, actual["Lead_Time"], target["lead_time"]))
        days_until = to_float(actual["Days_Until_Safety_Stock"])
        if days_until is None or abs(days_until - target["days_until"]) > 0.11:
            days_errors.append((product_id, actual["Days_Until_Safety_Stock"], target["days_until"]))
        needs = normalized(actual["Needs_Reorder"]) == "yes"
        if needs:
            actual_reorder_ids.add(product_id)
        if needs != target["needs_reorder"] or normalized(actual["Needs_Reorder"]) not in {"yes", "no"}:
            needs_errors.append((product_id, actual["Needs_Reorder"], target["needs_reorder"]))
        if target["needs_reorder"]:
            actual_date = parse_date(actual["Reorder_Date"])
            if actual_date != target["reorder_date"]:
                date_errors.append((product_id, str(actual["Reorder_Date"]), target["reorder_date"].isoformat()))
        elif normalized(actual["Reorder_Date"]) != "n/a":
            date_errors.append((product_id, str(actual["Reorder_Date"]), "N/A"))

    average_ok = history_days == 361 and eligible_order_count == 100 and not average_errors and actual_ids == expected_ids
    results.append(criterion("R003", average_ok, f"历史天数={history_days}，有效订单={eligible_order_count}，平均销量错误={average_errors[:8]}"))
    results.append(criterion("R004", not lead_errors and actual_ids == expected_ids, f"交付期错误={lead_errors[:8]}"))
    results.append(criterion("R005", not days_errors and actual_ids == expected_ids, f"可售天数错误={days_errors[:8]}"))

    expected_reorder_ids = {product_id for product_id, target in expected.items() if target["needs_reorder"]}
    needs_ok = not needs_errors and actual_reorder_ids == expected_reorder_ids and len(actual_reorder_ids) == 18
    results.append(criterion("R006", needs_ok, f"标记错误={needs_errors[:8]}，实际/期望补货={len(actual_reorder_ids)}/{len(expected_reorder_ids)}"))
    date_ok = not date_errors and actual_ids == expected_ids
    results.append(criterion("R007", date_ok, f"日期错误={date_errors[:8]}"))

    schedule_matched_ids = set()
    schedule_errors = []
    for record in schedule_records:
        matches = [
            product_id
            for product_id in expected_reorder_ids
            if name_matches(record["Product_Name"], expected[product_id]["name"])
        ]
        if len(matches) != 1:
            schedule_errors.append((record["Product_Name"], f"名称匹配数={len(matches)}"))
            continue
        product_id = matches[0]
        if product_id in schedule_matched_ids:
            schedule_errors.append((record["Product_Name"], "重复补货产品"))
            continue
        schedule_matched_ids.add(product_id)
        target = expected[product_id]
        fields_ok = (
            normalized(record["Category"]) == normalized(target["category"])
            and to_int(record["Current_Stock"]) == target["stock"]
            and parse_date(record["Reorder_Date"]) == target["reorder_date"]
            and to_int(record["Lead_Time_Days"]) == target["lead_time"]
        )
        if not fields_ok:
            schedule_errors.append((product_id, record))
    schedule_ok = (
        len(schedule_records) == 18
        and schedule_matched_ids == expected_reorder_ids
        and not schedule_errors
        and schedule_matched_ids == actual_reorder_ids
    )
    results.append(
        criterion(
            "R008",
            schedule_ok,
            f"行数={len(schedule_records)}/18，匹配={len(schedule_matched_ids)}，缺失ID={sorted(expected_reorder_ids - schedule_matched_ids)}，错误={schedule_errors[:6]}",
        )
    )

    roundtrip_ok = roundtrip_path is not None and structure_ok
    results.append(criterion("R009", roundtrip_ok, roundtrip_evidence))
    results.append(criterion("R010", not source_ok, f"来源覆盖通过={source_ok}"))
    reorder_error = not (needs_ok and date_ok and schedule_ok)
    results.append(criterion("R011", reorder_error, f"补货标记={needs_ok}，日期={date_ok}，日程表={schedule_ok}"))
    results.append(criterion("R012", not roundtrip_ok or not structure_ok, f"结构={structure_ok}，{roundtrip_evidence}"))
    return build_result(results, output_path)


def build_result(results: list[dict], artifact_path: Path) -> dict:
    present = {item["criterion_id"] for item in results}
    for identifier in CRITERION_SPECS:
        if identifier not in present:
            results.append(criterion(identifier, False, "因交付不存在、损坏或门槛失败，本判据未执行"))
    results.sort(key=lambda item: item["criterion_id"])

    positive_total = sum(item["weight"] for item in results if item["weight"] > 0)
    earned = sum(item["weight"] for item in results if item["weight"] > 0 and item["passed"])
    penalties = sum(item["weight"] for item in results if item["weight"] < 0 and item["passed"])
    reward = max(0.0, min(1.0, (earned + penalties) / positive_total))
    hurdle_failures = [item["criterion_id"] for item in results if item["type"] == "hurdle" and not item["passed"]]
    if hurdle_failures:
        reward = min(reward, min(HURDLE_CAPS[identifier] for identifier in hurdle_failures))
    digest = hashlib.sha256(artifact_path.read_bytes()).hexdigest() if artifact_path.is_file() else hashlib.sha256(b"").hexdigest()

    criterion_results = []
    for item in results:
        score = (0.0 if item["passed"] else 1.0) if item["type"] == "penalty" else item["score"]
        criterion_results.append(
            {
                "criterion_id": item["criterion_id"],
                "score": score,
                "evidence": item["evidence"],
                "method": item["method"],
                "raw": {
                    "description": item["description"],
                    "weight": item["weight"],
                    "type": item["type"],
                    "passed": item["passed"],
                },
            }
        )
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [digest],
        "criterion_results": criterion_results,
        "reward": round(reward, 6),
        "passed": reward >= 0.75,
        "hurdle_failures": hurdle_failures,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path)
    parser.add_argument("--reward-file", type=Path)
    args = parser.parse_args()
    result = evaluate(args.artifact_dir)
    rendered = json.dumps(result, ensure_ascii=False, indent=2)
    print(rendered)
    if args.result:
        args.result.parent.mkdir(parents=True, exist_ok=True)
        args.result.write_text(rendered + "\n", encoding="utf-8")
    if args.reward_file:
        args.reward_file.parent.mkdir(parents=True, exist_ok=True)
        args.reward_file.write_text(f"{result['reward']:.6f}\n", encoding="utf-8")
    raise SystemExit(0 if result["passed"] else 1)


if __name__ == "__main__":
    main()
