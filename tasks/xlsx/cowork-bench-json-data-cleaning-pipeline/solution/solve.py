#!/usr/bin/env python3
"""从两份 JSON 生成唯一的可审计清洗工作簿。"""
from __future__ import annotations
import argparse, sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from clean_data import process
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

def style(sheet, widths, ref, name):
    for cell in sheet[1]:
        cell.fill=PatternFill("solid",fgColor="17365D"); cell.font=Font(name="Arial",size=10,bold=True,color="FFFFFF"); cell.alignment=Alignment(horizontal="center",wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row: cell.border=Border(bottom=Side(style="thin",color="D7E0EA")); cell.alignment=Alignment(vertical="center",wrap_text=True)
    for col,width in widths.items(): sheet.column_dimensions[col].width=width
    sheet.freeze_panes="A2"; sheet.auto_filter.ref=ref; sheet.sheet_view.showGridLines=False
    table=Table(displayName=name,ref=ref); table.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True); sheet.add_table(table)

def build(input_dir: Path, output_dir: Path) -> None:
    _, _, result=process(input_dir/"sales_api_response.json",input_dir/"inventory_api_response.json"); wb=Workbook(); wb.calculation.fullCalcOnLoad=True; wb.calculation.forceFullCalc=True; wb.calculation.calcMode="auto"
    sales=wb.active; sales.title="Sales Analysis"; headers=["Product_Name","Total_Quantity_Sold","Total_Revenue","Avg_Unit_Price"]; sales.append(headers)
    for row in result["sales_analysis"]: sales.append([row[k] for k in headers])
    for col in ("C","D"):
        for cell in sales[col][1:]: cell.number_format='₽#,##0.00'
    style(sales,{"A":28,"B":23,"C":20,"D":20},f"A1:D{sales.max_row}","Pipeline_Sales_Analysis")
    inventory=wb.create_sheet("Inventory Status"); inventory.append(["Product_Name","Current_Stock","Reorder_Point","Status"])
    for i,row in enumerate(result["inventory_status"],2): inventory.append([row["Product_Name"],row["Current_Stock"],row["Reorder_Point"],f'=IF(B{i}<=C{i},"LOW","OK")'])
    style(inventory,{"A":28,"B":20,"C":20,"D":16},f"A1:D{inventory.max_row}","Pipeline_Inventory_Status")
    combined=wb.create_sheet("Combined View"); combined_headers=["Product_Name","Total_Sold","Current_Stock","Revenue","Stock_Status"]; combined.append(combined_headers)
    for row in result["combined_view"]: combined.append([row[k] for k in combined_headers])
    for cell in combined["D"][1:]: cell.number_format='₽#,##0.00'
    style(combined,{"A":28,"B":17,"C":20,"D":18,"E":18},f"A1:E{combined.max_row}","Pipeline_Combined_View")
    dashboard=wb.create_sheet("Pipeline Dashboard"); dashboard.append(combined_headers)
    for row in result["combined_view"]: dashboard.append([row[k] for k in combined_headers])
    for cell in dashboard["D"][1:]: cell.number_format='₽#,##0.00'
    style(dashboard,{"A":28,"B":17,"C":20,"D":18,"E":18},f"A1:E{dashboard.max_row}","Pipeline_Dashboard")
    summary=wb.create_sheet("Run Summary"); summary.append(["Metric","Value"])
    for key,value in result["summary"].items(): summary.append([key,value])
    style(summary,{"A":34,"B":24},f"A1:B{summary.max_row}","Pipeline_Run_Summary")
    output_dir.mkdir(parents=True,exist_ok=True); wb.save(output_dir/"Data_Pipeline_Report.xlsx"); wb.close()

def main():
    ap=argparse.ArgumentParser(); ap.add_argument("--input-dir",type=Path,default=Path("/app/data/input_files")); ap.add_argument("--output-dir",type=Path,default=Path("/app/output")); args=ap.parse_args(); build(args.input_dir,args.output_dir)
if __name__=="__main__": main()
