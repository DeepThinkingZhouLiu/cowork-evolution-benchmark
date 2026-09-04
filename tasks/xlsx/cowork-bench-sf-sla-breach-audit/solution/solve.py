#!/usr/bin/env python3
"""Cowork-Bench SLA 违约审计题的逐题参考解法。"""

from __future__ import annotations

import argparse
import csv
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


WORKBOOK_NAME = "SLA_Breach_Audit.xlsx"
ANALYSIS_HEADERS = [
    "Priority",
    "Total_Tickets",
    "Avg_Response_Hours",
    "SLA_Target_Hours",
    "Breached_Count",
    "Breach_Rate",
]
SUMMARY_METRICS = [
    "Total_Tickets",
    "Total_Breached",
    "Overall_Breach_Rate",
    "Worst_Priority",
    "Best_Priority",
]


def round_2(value: Decimal) -> Decimal:
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def calculate(input_dir: Path) -> list[dict]:
    with (input_dir / "sla_policies.csv").open(encoding="utf-8", newline="") as handle:
        policies = {
            row["PRIORITY"]: Decimal(row["RESPONSE_TARGET_HOURS"])
            for row in csv.DictReader(handle)
        }

    aggregates = defaultdict(lambda: {"count": 0, "response_sum": Decimal("0"), "breached": 0})
    with (input_dir / "support_tickets.csv").open(encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle):
            priority = row["PRIORITY"]
            if priority not in policies:
                raise ValueError(f"工单优先级缺少 SLA 策略：{priority}")
            response = Decimal(row["RESPONSE_TIME_HOURS"])
            aggregate = aggregates[priority]
            aggregate["count"] += 1
            aggregate["response_sum"] += response
            if response > policies[priority]:
                aggregate["breached"] += 1

    rows = []
    for priority in sorted(aggregates, key=lambda item: (policies[item], item)):
        aggregate = aggregates[priority]
        count = aggregate["count"]
        rows.append(
            {
                "priority": priority,
                "total": count,
                "average": round_2(aggregate["response_sum"] / count),
                "target": policies[priority],
                "breached": aggregate["breached"],
                "rate": round_2(Decimal(aggregate["breached"]) * 100 / count),
            }
        )
    return rows


def add_table(sheet, ref: str, name: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def style_sheet(sheet, widths: dict[str, float]) -> None:
    navy = "17365D"
    white = "FFFFFF"
    border_color = "D9E2F3"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 26
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=11, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color=border_color)
    for row in sheet.iter_rows():
        for cell in row:
            cell.font = Font(name="Aptos", size=10, bold=cell.row == 1, color=white if cell.row == 1 else "1F2937")
            cell.alignment = Alignment(vertical="center", horizontal="left" if cell.column == 1 else "right")
            cell.border = Border(bottom=thin)
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_options.horizontalCentered = True
    sheet.oddFooter.center.text = "Cowork-Bench 静态快照 · SLA Breach Audit"


def write_workbook(output_path: Path, rows: list[dict]) -> None:
    workbook = Workbook()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.properties.creator = "Cowork Evolution Benchmark"
    workbook.properties.title = "SLA Breach Audit"
    workbook.properties.description = "从 Cowork-Bench SUPPORT_CENTER 静态快照审计首次响应 SLA。"

    analysis = workbook.active
    analysis.title = "Breach Analysis"
    analysis.append(ANALYSIS_HEADERS)
    for position, row in enumerate(rows, start=2):
        analysis.append(
            [
                row["priority"],
                row["total"],
                float(row["average"]),
                float(row["target"]),
                row["breached"],
                f"=ROUND(E{position}/B{position}*100,2)",
            ]
        )
    analysis["A1"].comment = Comment("PRIORITY；只报告实际存在工单的优先级。", "Cowork Evolution Benchmark")
    analysis["C1"].comment = Comment("RESPONSE_TIME_HOURS 的算术平均值，两位小数。", "Cowork Evolution Benchmark")
    analysis["D1"].comment = Comment("sla_policies.csv 中的 RESPONSE_TARGET_HOURS。", "Cowork Evolution Benchmark")
    analysis["E1"].comment = Comment("违约定义：RESPONSE_TIME_HOURS 严格大于 SLA 目标。", "Cowork Evolution Benchmark")
    analysis["F1"].comment = Comment("Breached_Count / Total_Tickets × 100。", "Cowork Evolution Benchmark")
    for cell in analysis["B"][1:]:
        cell.number_format = "#,##0"
    for column in ("C", "D"):
        for cell in analysis[column][1:]:
            cell.number_format = "0.00"
    for cell in analysis["E"][1:]:
        cell.number_format = "#,##0"
    for cell in analysis["F"][1:]:
        cell.number_format = '0.00"%"'
    style_sheet(analysis, {"A": 18, "B": 16, "C": 23, "D": 21, "E": 19, "F": 16})
    add_table(analysis, f"A1:F{analysis.max_row}", "SLA_Breach_Analysis")
    analysis.conditional_formatting.add(
        f"F2:F{analysis.max_row}",
        ColorScaleRule(start_type="min", start_color="63BE7B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="F8696B"),
    )

    summary = workbook.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    summary.append([SUMMARY_METRICS[0], "=SUM('Breach Analysis'!B2:B4)"])
    summary.append([SUMMARY_METRICS[1], "=SUM('Breach Analysis'!E2:E4)"])
    summary.append([SUMMARY_METRICS[2], "=ROUND(B3/B2*100,2)"])
    summary.append([SUMMARY_METRICS[3], "=INDEX('Breach Analysis'!A2:A4,MATCH(MAX('Breach Analysis'!F2:F4),'Breach Analysis'!F2:F4,0))"])
    summary.append([SUMMARY_METRICS[4], "=INDEX('Breach Analysis'!A2:A4,MATCH(MIN('Breach Analysis'!F2:F4),'Breach Analysis'!F2:F4,0))"])
    for cell in summary["B"][1:3]:
        cell.number_format = "#,##0"
    summary["B4"].number_format = '0.00"%"'
    style_sheet(summary, {"A": 28, "B": 22})
    add_table(summary, f"A1:B{summary.max_row}", "SLA_Audit_Summary")

    total = sum(row["total"] for row in rows)
    breached = sum(row["breached"] for row in rows)
    overall_rate = round_2(Decimal(breached) * 100 / total)
    worst = max(rows, key=lambda row: row["rate"])["priority"]
    draft = workbook.create_sheet("Escalation Draft")
    draft.append(["Field", "Value"])
    draft.append(["From", "sla-audit@company.com"])
    draft.append(["To", "ops-director@company.com"])
    draft.append(["Subject", "SLA Breach Audit Results"])
    draft.append(["Body", "The SLA breach audit is complete."])
    draft.append(["Overall_Breach_Rate", float(overall_rate)])
    draft.append(["Worst_Priority", worst])
    draft.append(["Attachment", WORKBOOK_NAME])
    style_sheet(draft, {"A": 24, "B": 72})
    draft["B6"].number_format = '0.00"%"'

    workbook.save(output_path)


def solve(input_dir: Path, output_dir: Path) -> None:
    rows = calculate(input_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    write_workbook(output_dir / WORKBOOK_NAME, rows)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=Path("/app/data/input_files"))
    parser.add_argument("--output-dir", type=Path, default=Path("/app/output"))
    args = parser.parse_args()
    solve(args.input_dir, args.output_dir)


if __name__ == "__main__":
    main()
