#!/usr/bin/env python3
"""从 CSV 快照独立重算 SLA 分组结果，检查单一工作簿交付。"""
from __future__ import annotations
import argparse, csv, hashlib, json, zipfile
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from openpyxl import load_workbook

TASK_ID="cowork-evo/cowork-bench-sf-sla-breach-audit"; RUBRIC_VERSION="cowork-bench-sf-sla-breach-audit-static-v0.2"; JUDGE_VERSION="cowork-bench-sf-sla-breach-audit-static-v0.2"; OUTPUT="SLA_Breach_Audit.xlsx"
HEADERS=["Priority","Total_Tickets","Avg_Response_Hours","SLA_Target_Hours","Breached_Count","Breach_Rate"]; WEIGHTS={"R001":.10,"R002":.20,"R003":.25,"R004":.20,"R005":.10,"R006":.10,"R007":.05}
def criterion(cid,ok,evidence): return {"criterion_id":cid,"score":float(ok),"evidence":evidence,"method":"deterministic"}
def digest(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def r2(v): return Decimal(v).quantize(Decimal("0.01"),rounding=ROUND_HALF_UP)
def main():
    ap=argparse.ArgumentParser(); ap.add_argument("--artifact-dir",type=Path,required=True); ap.add_argument("--result",type=Path,required=True); ap.add_argument("--reward-file",type=Path,required=True); args=ap.parse_args(); root=Path(__file__).parents[1]; inp=root/"data/input_files"; out=args.artifact_dir/OUTPUT; results=[]
    try:
        with (inp/"sla_policies.csv").open(encoding="utf-8",newline="") as f: policies={r["PRIORITY"]:Decimal(r["RESPONSE_TARGET_HOURS"]) for r in csv.DictReader(f)}
        agg=defaultdict(lambda:{"n":0,"sum":Decimal("0"),"breach":0})
        with (inp/"support_tickets.csv").open(encoding="utf-8",newline="") as f:
            for row in csv.DictReader(f):
                p=row["PRIORITY"]; response=Decimal(row["RESPONSE_TIME_HOURS"]); agg[p]["n"]+=1; agg[p]["sum"]+=response; agg[p]["breach"]+=int(response>policies[p])
        expected=[]
        for p in sorted(agg,key=lambda x:(policies[x],x)): expected.append((p,agg[p]["n"],r2(agg[p]["sum"]/agg[p]["n"]),policies[p],agg[p]["breach"],r2(Decimal(agg[p]["breach"])*100/agg[p]["n"])))
        total=sum(x[1] for x in expected); breached=sum(x[4] for x in expected); overall=r2(Decimal(breached)*100/total); worst=max(expected,key=lambda x:x[5])[0]; best=min(expected,key=lambda x:x[5])[0]
        valid=out.is_file() and zipfile.is_zipfile(out)
        if not valid:
            results=[criterion("R001",False,"输出缺失或损坏")]+[criterion(cid,False,"缺少有效输出") for cid in ("R002","R003","R004","R005","R006")]+[criterion("R007",True,"输入快照固定")]
        else:
            wb=load_workbook(out,data_only=False); structure=wb.sheetnames==["Breach Analysis","Summary","Escalation Draft"]; results.append(criterion("R001",structure,f"工作表={wb.sheetnames}")); analysis=wb["Breach Analysis"]; rows=list(analysis.iter_rows(min_row=2,max_col=6,values_only=False)); actual=[]
            for row in rows:
                if row[0].value is not None: actual.append(tuple([row[0].value,row[1].value,row[2].value,row[3].value,row[4].value,row[5].value]))
            data_ok=len(actual)==len(expected) and tuple(x.value for x in analysis[1][:6])==tuple(HEADERS) and all(str(a[0])==str(e[0]) and a[1]==e[1] and abs(float(a[2])-float(e[2]))<=.011 and abs(float(a[3])-float(e[3]))<=.001 and a[4]==e[4] for a,e in zip(actual,expected)); results.append(criterion("R002",data_ok,"优先级、数量、均值和目标检查")); rate_ok=data_ok and all(abs(float(a[4])*100/float(a[1])-float(e[5]))<=.011 and isinstance(a[5],str) and a[5].startswith("=") for a,e in zip(actual,expected)); results.append(criterion("R003",rate_ok,"违约数/率及公式检查"))
            ss=wb["Summary"]; srows=list(ss.iter_rows(min_row=2,max_col=2,values_only=True)); expected_labels=["Total_Tickets","Total_Breached","Overall_Breach_Rate","Worst_Priority","Best_Priority"]; labels_ok=[str(row[0]) for row in srows]==expected_labels; formula_cells=[ss[f"B{i}"].value for i in range(2,7)]; summary_ok=len(srows)==5 and labels_ok and all(isinstance(v,str) and v.startswith("=") for v in formula_cells) and "SUM" in formula_cells[0].upper() and "SUM" in formula_cells[1].upper() and "ROUND" in formula_cells[2].upper() and "MAX" in formula_cells[3].upper() and "MIN" in formula_cells[4].upper(); results.append(criterion("R004",summary_ok,"Summary 标签和公式检查"))
            draft=wb["Escalation Draft"]; text=" ".join(str(v or "") for row in draft.iter_rows(values_only=True) for v in row); draft_ok=all(x in text for x in ("sla-audit@company.com","ops-director@company.com","SLA Breach Audit Results","Overall_Breach_Rate","Worst_Priority",worst)); results.append(criterion("R005",draft_ok,"升级草稿字段和结论检查")); professional=all(wb[s].freeze_panes=="A2" and wb[s].auto_filter.ref and wb[s].protection.sheet is not True for s in wb.sheetnames); results.append(criterion("R006",professional,"筛选、冻结、原生可编辑检查")); results.append(criterion("R007",True,"输入快照固定且 provenance 已记录")); wb.close()
    except Exception as exc:
        results=[criterion("R001",False,f"读取失败：{exc}")]+[criterion(cid,False,"读取失败") for cid in ("R002","R003","R004","R005","R006","R007")]
    vals={x["criterion_id"]:x["score"] for x in results}; reward=sum(WEIGHTS[k]*vals.get(k,0) for k in WEIGHTS); passed=all(vals.get(k,0)==1 for k in ("R001","R002","R003","R004","R005")) and reward>=.8; payload={"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":[digest(out)] if out.is_file() else [],"criterion_results":results,"reward":reward,"passed":passed}; args.result.parent.mkdir(parents=True,exist_ok=True); args.result.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8"); args.reward_file.parent.mkdir(parents=True,exist_ok=True); args.reward_file.write_text(f"{reward:.6f}\n"); raise SystemExit(0 if passed else 1)
if __name__=="__main__": main()
