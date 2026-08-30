#!/usr/bin/env python3
"""供应商合同合规审计的逐题参考解法。

所有业务状态都从随题 CSV/PDF 输入计算；外部邮件和日历服务被输出草稿替代。
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import re
import subprocess
from datetime import date, timedelta
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

INPUT_CSV = "supplier_list.csv"
TEMPLATE = "contract_template.pdf"
EXPECTED_INPUTS = {
    "supplier_list.csv": "4553ce4cfd4ce8c336c900f433e1530d07c82b404a71eaa7e157ba9b71c78a1f",
    "contract_template.pdf": "f0b935f18a9a05772eb7b7f6bdaf98f7999ce9e55d50f4ee66452695e118df06",
    "contract_supplier1_TechService.pdf": "b54b1d5bc3a5e7bf459718e0a0aa2f337917a3646a1abb1f6ab725928f6edcc1",
    "contract_supplier2_GlobalLogistics.pdf": "b3f3fc8f20ed819c651dcb4435407680ea1da78785ccdaef02ab1d85cdddd98d",
    "contract_supplier3_ControlQuality.pdf": "066dfadfd2cf5a5e0d59e7e264fc1b0e0b4f5f7b8de50c38790c2646b1221fe8",
    "contract_supplier4_ConsultGroup.pdf": "9fbc24147b38bedca6202bf82ba2ae04df4ec098af0ccf6e6f2b142774ed8843",
}
EXPECTED_INPUTS["contract_supplier5_ServicePodderzhka.pdf"] = "b7423c3e062654c76a1395685426e3a230be02fc7f2ede865dbc9ba515cd9907"

CONTRACT_FILES = [
    "contract_supplier1_TechService.pdf",
    "contract_supplier2_GlobalLogistics.pdf",
    "contract_supplier3_ControlQuality.pdf",
    "contract_supplier4_ConsultGroup.pdf",
    "contract_supplier5_ServicePodderzhka.pdf",
]
DIMENSIONS = ["Scope Clarity", "SLA Present", "Payment Terms", "IP Rights", "Confidentiality", "Liability Limits"]
SUPPLIER_LABELS = {
    "Scope Clarity": "Scope Clarity",
    "SLA Present": "SLA Present",
    "Payment Terms": "Payment Terms",
    "IP Rights": "IP Rights",
    "Confidentiality": "Confidentiality",
    "Liability Limits": "Liability Limits",
}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def pdf_text(path: Path) -> str:
    result = subprocess.run(["pdftotext", "-layout", str(path), "-"], check=True, capture_output=True, text=True)
    return " ".join(result.stdout.lower().split())


def audit_contract(text: str) -> dict[str, str]:
    low = text
    scope = "Incomplete" if "не детализированы" in low else "Complete"
    sla = "Complete" if all(token in low for token in ["99,5%", "4 часа", "24 часа"]) else "Incomplete"
    payment = "Compliant" if "в течение 30 дней" in low and "в течение 90 дней" not in low else "Non-Compliant"
    ip = "Compliant" if "все результаты работ" in low and "собственностью компании" in low else "Incomplete"
    confidentiality = "Compliant" if "конфиденциальность сведений" in low else "Missing"
    liability = "Compliant" if "совокупная ответственность ограничена" in low else "Missing"
    return {
        "Scope Clarity": scope,
        "SLA Present": sla,
        "Payment Terms": payment,
        "IP Rights": ip,
        "Confidentiality": confidentiality,
        "Liability Limits": liability,
    }


def is_good(dimension: str, status: str) -> bool:
    return status == ("Complete" if dimension in {"Scope Clarity", "SLA Present", "IP Rights", "Confidentiality"} else "Compliant")


def build_records(input_dir: Path) -> list[dict]:
    for filename, expected in EXPECTED_INPUTS.items():
        actual = sha256(input_dir / filename)
        if actual != expected:
            raise ValueError(f"输入哈希不匹配：{filename} {actual}")
    with (input_dir / INPUT_CSV).open(encoding="utf-8-sig", newline="") as handle:
        suppliers = list(csv.DictReader(handle))
    if len(suppliers) != 5:
        raise ValueError(f"供应商行数异常：{len(suppliers)}")
    records = []
    for row, filename in zip(suppliers, CONTRACT_FILES):
        statuses = audit_contract(pdf_text(input_dir / filename))
        good_count = sum(1 for dim, status in statuses.items() if (status == "Complete" if dim in {"Scope Clarity", "SLA Present"} else status == "Compliant"))
        gaps = [dim for dim, status in statuses.items() if not ((status == "Complete" if dim in {"Scope Clarity", "SLA Present"} else status == "Compliant"))]
        score = round(good_count * 100 / len(DIMENSIONS), 2)
        risk = "High" if len(gaps) >= 2 else "Medium" if gaps else "Low"
        records.append({"supplier": row["Supplier_Name"], "email": row["Contact_Email"], "service": row["Service_Type"], "start": row["Contract_Start_Date"], "annual_value": row["Annual_Contract_Value"], "statuses": statuses, "gaps": gaps, "score": score, "risk": risk})
    return records


def style_sheet(ws, table_name: str, end_row: int, end_col: int) -> None:
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2, max_row=end_row, max_col=end_col):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    tab = Table(displayName=table_name, ref=f"A1:{chr(64 + end_col)}{end_row}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(tab)
    for column in ws.columns:
        letter = column[0].column_letter
        ws.column_dimensions[letter].width = min(42, max(14, max(len(str(c.value or "")) for c in column) + 2))


def write_workbooks(records: list[dict], output_dir: Path) -> None:
    wb = Workbook(); ws = wb.active; ws.title = "Compliance Assessment"
    headers = ["Supplier", *DIMENSIONS, "Compliance Score"]
    ws.append(headers)
    for rec in records:
        ws.append([rec["supplier"], *[rec["statuses"][dim] for dim in DIMENSIONS], rec["score"]])
    style_sheet(ws, "ComplianceAssessment", ws.max_row, len(headers))
    ws.freeze_panes = "A2"
    wb.save(output_dir / "compliance_assessment.xlsx")

    wb = Workbook(); ws = wb.active; ws.title = "Corrective Actions"
    headers = ["Supplier", "Issue", "Required Action", "Owner", "Target Date", "Est. Cost", "Status"]
    ws.append(headers)
    owners = ["采购合规负责人", "法务负责人"]
    action_templates = {
        "Scope Clarity": ("补充服务范围、交付成果和期限", "明确服务范围并签署补充条款"),
        "SLA Present": ("合同未给出具体 SLA 指标", "补充可衡量的可用性、响应和修复指标"),
        "Payment Terms": ("付款条件偏离标准 30 天条款", "由采购与财务复核并修订付款条款"),
        "IP Rights": ("成果知识产权归属不明确", "明确公司对合同成果的权利"),
        "Confidentiality": ("合同缺少保密条款", "补充保密义务和期限"),
        "Liability Limits": ("合同缺少责任上限条款", "补充不超过 12 个月付款额的责任上限"),
    }
    cost_by_dim = {"Scope Clarity": 0, "SLA Present": 370000, "Payment Terms": 740000, "IP Rights": 150000, "Confidentiality": 0, "Liability Limits": 0}
    for idx, rec in enumerate(records):
        for gap_idx, dim in enumerate(rec["gaps"]):
            issue, action = action_templates[dim]
            target = date.today() + timedelta(days=21 + idx * 7 + gap_idx * 7)
            ws.append([rec["supplier"], issue, action, owners[gap_idx % len(owners)], target.isoformat(), cost_by_dim[dim], "Planned"])
    style_sheet(ws, "CorrectiveActions", ws.max_row, len(headers))
    wb.save(output_dir / "corrective_actions.xlsx")


def add_doc_table(doc: Document, headers: list[str], rows: list[list[object]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers)); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]; cell.text = header
        for run in cell.paragraphs[0].runs: run.font.bold = True
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values): cells[i].text = str(value)


def write_docs(records: list[dict], output_dir: Path) -> None:
    findings = Document(); findings.add_heading("供应商合同合规审计发现", 0)
    findings.add_paragraph("本报告仅根据供应商清单、合同模板和五份合同 PDF 的文字内容编制。")
    for rec in records:
        findings.add_heading(rec["supplier"], 1)
        findings.add_paragraph(f"服务类型：{rec['service']}；合同开始日期：{rec['start']}；年度合同金额：{rec['annual_value']}。")
        if rec["gaps"]:
            findings.add_paragraph("具体差距：" + "；".join(f"{dim}={rec['statuses'][dim]}" for dim in rec["gaps"]) + "。")
            findings.add_paragraph("建议：" + "；".join({"Scope Clarity": "明确服务范围", "SLA Present": "补充可衡量 SLA", "Payment Terms": "复核付款条件", "IP Rights": "明确知识产权归属", "Confidentiality": "补充保密条款", "Liability Limits": "补充责任上限"}[dim] for dim in rec["gaps"]) + "。")
        else:
            findings.add_paragraph("未发现与模板六个审计维度相关的差距，建议按现有条款持续监控。")
        findings.add_paragraph(f"Compliance Score: {rec['score']:.2f} / 100；Overall Risk: {rec['risk']}。")
    findings.save(output_dir / "audit_findings.docx")

    summary = Document(); summary.add_heading("合同合规审计执行摘要", 0)
    average = sum(rec["score"] for rec in records) / len(records)
    counts = {risk: sum(rec["risk"] == risk for rec in records) for risk in ("High", "Medium", "Low")}
    summary.add_paragraph(f"Average Compliance Score: {average:.2f} / 100。")
    summary.add_paragraph(f"Risk distribution: High={counts['High']}; Medium={counts['Medium']}; Low={counts['Low']}。")
    add_doc_table(summary, ["Supplier", "Compliance Score", "Overall Risk"], [[rec["supplier"], f"{rec['score']:.2f}", rec["risk"]] for rec in records])
    summary.add_paragraph("高风险供应商需要优先完成整改；本文件和相关邮件/会议文件均为离线草稿，未执行外部发送或日历创建。")
    summary.save(output_dir / "audit_summary.docx")


def write_email(records: list[dict], output_dir: Path) -> None:
    recipients = ", ".join(rec["email"] for rec in records)
    lines = ["From: procurement-compliance@company.invalid", f"To: {recipients}", "Subject: Contract Compliance Audit Findings", "Content-Type: text/plain; charset=utf-8", "", "合同合规审计通知（离线草稿，未发送）", ""]
    for rec in records:
        lines.append(f"{rec['supplier']}：Compliance Score {rec['score']:.2f}/100，风险 {rec['risk']}。")
        if rec["gaps"]: lines.append("整改事项：" + "、".join(rec["gaps"]) + "。")
        else: lines.append("未发现需要整改的模板维度。")
    (output_dir / "audit_notifications.eml").write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_calendar(records: list[dict], output_dir: Path) -> None:
    high = [rec for rec in records if rec["risk"] == "High"]
    lines = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//CoworkEvoBench//Supplier Compliance//EN", "CALSCALE:GREGORIAN", "METHOD:PUBLISH"]
    for idx, rec in enumerate(high):
        day = date.today() + timedelta(days=30 + idx)
        start = day.strftime("%Y%m%d") + "T090000"
        end = day.strftime("%Y%m%d") + "T100000"
        lines += ["BEGIN:VEVENT", f"UID:supplier-review-{idx}@coworkevobench.local", f"DTSTAMP:{date.today().strftime('%Y%m%d')}T000000Z", f"DTSTART:{start}", f"DTEND:{end}", f"SUMMARY:Supplier compliance improvement review - {rec['supplier']}", f"DESCRIPTION:Owner: procurement compliance lead; Corrective actions: {', '.join(rec['gaps'])}", "END:VEVENT"]
    lines += ["END:VCALENDAR", ""]
    (output_dir / "supplier_review_meetings.ics").write_text("\r\n".join(lines), encoding="utf-8", newline="")


def main() -> None:
    parser = argparse.ArgumentParser(); parser.add_argument("--input-dir", type=Path, required=True); parser.add_argument("--output-dir", type=Path, required=True); args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    records = build_records(args.input_dir)
    write_workbooks(records, args.output_dir)
    write_docs(records, args.output_dir)
    write_email(records, args.output_dir)
    write_calendar(records, args.output_dir)


if __name__ == "__main__": main()
