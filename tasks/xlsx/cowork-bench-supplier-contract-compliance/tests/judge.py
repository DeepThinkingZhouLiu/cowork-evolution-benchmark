#!/usr/bin/env python3
"""确定性离线 Judge：从 CSV/PDF 输入独立重算合同合规事实。"""

from __future__ import annotations

import argparse
import csv
import hashlib
import re
import subprocess
import zipfile
from datetime import date
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

TASK_ID = "cowork-evo/cowork-bench-supplier-contract-compliance"
RUBRIC_VERSION = "cowork-bench-supplier-contract-compliance-v0.1"
JUDGE_VERSION = "cowork-bench-supplier-contract-compliance-deterministic-v0.1"
OUTPUTS = ["compliance_assessment.xlsx", "corrective_actions.xlsx", "audit_findings.docx", "audit_summary.docx", "audit_notifications.eml", "supplier_review_meetings.ics"]
CONTRACTS = ["contract_supplier1_TechService.pdf", "contract_supplier2_GlobalLogistics.pdf", "contract_supplier3_ControlQuality.pdf", "contract_supplier4_ConsultGroup.pdf", "contract_supplier5_ServicePodderzhka.pdf"]
DIMENSIONS = ["Scope Clarity", "SLA Present", "Payment Terms", "IP Rights", "Confidentiality", "Liability Limits"]
EXPECTED_HASHES = {
    "supplier_list.csv": "4553ce4cfd4ce8c336c900f433e1530d07c82b404a71eaa7e157ba9b71c78a1f",
    "contract_template.pdf": "f0b935f18a9a05772eb7b7f6bdaf98f7999ce9e55d50f4ee66452695e118df06",
    "contract_supplier1_TechService.pdf": "b54b1d5bc3a5e7bf459718e0a0aa2f337917a3646a1abb1f6ab725928f6edcc1",
    "contract_supplier2_GlobalLogistics.pdf": "b3f3fc8f20ed819c651dcb4435407680ea1da78785ccdaef02ab1d85cdddd98d",
    "contract_supplier3_ControlQuality.pdf": "066dfadfd2cf5a5e0d59e7e264fc1b0e0b4f5f7b8de50c38790c2646b1221fe8",
    "contract_supplier4_ConsultGroup.pdf": "9fbc24147b38bedca6202bf82ba2ae04df4ec098af0ccf6e6f2b142774ed8843",
    "contract_supplier5_ServicePodderzhka.pdf": "b7423c3e062654c76a1395685426e3a230be02fc7f2ede865dbc9ba515cd9907",
}


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def pdf_text(path: Path) -> str:
    raw = subprocess.run(["pdftotext", "-layout", str(path), "-"], check=True, capture_output=True, text=True).stdout.lower()
    return " ".join(raw.split())


def expected_records(source_dir: Path) -> list[dict]:
    with (source_dir / "supplier_list.csv").open(encoding="utf-8-sig", newline="") as handle:
        suppliers = list(csv.DictReader(handle))
    template = pdf_text(source_dir / "contract_template.pdf")
    if "соглашение об уровне обслуживания" not in template or "в течение 30 дней" not in template:
        raise ValueError("模板缺少标准合同证据")
    records = []
    for row, filename in zip(suppliers, CONTRACTS):
        text = pdf_text(source_dir / filename)
        statuses = {
            "Scope Clarity": "Incomplete" if "не детализированы" in text else "Complete",
            "SLA Present": "Complete" if all(x in text for x in ("99,5%", "4 часа", "24 часа")) else "Incomplete",
            "Payment Terms": "Compliant" if "в течение 30 дней" in text and "в течение 90 дней" not in text else "Non-Compliant",
            "IP Rights": "Compliant" if "все результаты работ" in text and "собственностью компании" in text else "Incomplete",
            "Confidentiality": "Compliant" if "конфиденциальность сведений" in text else "Missing",
            "Liability Limits": "Compliant" if "совокупная ответственность ограничена" in text else "Missing",
        }
        good = sum(status == ("Complete" if dim in {"Scope Clarity", "SLA Present"} else "Compliant") for dim, status in statuses.items())
        gaps = [dim for dim, status in statuses.items() if status != ("Complete" if dim in {"Scope Clarity", "SLA Present"} else "Compliant")]
        records.append({"supplier": row["Supplier_Name"], "email": row["Contact_Email"], "statuses": statuses, "gaps": gaps, "score": round(good * 100 / 6, 2), "risk": "High" if len(gaps) >= 2 else "Medium" if gaps else "Low"})
    return records


def criterion(cid: str, score: bool | float, evidence: str) -> dict:
    return {"criterion_id": cid, "score": float(score), "evidence": evidence, "method": "deterministic"}


def sheet_rows(path: Path, title: str) -> list[tuple]:
    wb = load_workbook(path, data_only=True)
    ws = next((wb[name] for name in wb.sheetnames if name.strip().lower() == title.lower()), None)
    if ws is None:
        raise ValueError(f"缺少工作表 {title}")
    rows = list(ws.iter_rows(values_only=True)); wb.close(); return rows


def doc_text(path: Path) -> str:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def parse_ics(path: Path) -> list[dict[str, str]]:
    text = path.read_text(encoding="utf-8")
    if "BEGIN:VCALENDAR" not in text or "END:VCALENDAR" not in text:
        raise ValueError("不是 VCALENDAR")
    events = []
    for body in re.findall(r"BEGIN:VEVENT\r?\n(.*?)\r?\nEND:VEVENT", text, flags=re.S):
        props = {}
        for line in body.splitlines():
            if ":" in line:
                key, value = line.split(":", 1); props[key.split(";", 1)[0].upper()] = value
        if not {"UID", "DTSTART", "DTEND", "SUMMARY", "DESCRIPTION"}.issubset(props):
            raise ValueError("VEVENT 缺少必需属性")
        events.append(props)
    return events


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    source_dir = Path(__file__).parents[1] / "data/input_files"
    results = []
    try:
        records = expected_records(source_dir)
        missing = [name for name in OUTPUTS if not (args.artifact_dir / name).is_file()]
        results.append(criterion("R001", not missing, f"缺失={missing}"))

        compliance_rows = sheet_rows(args.artifact_dir / "compliance_assessment.xlsx", "Compliance Assessment") if not missing or (args.artifact_dir / OUTPUTS[0]).is_file() else []
        header_idx = next((i for i, row in enumerate(compliance_rows) if set(DIMENSIONS + ["Supplier", "Compliance Score"]).issubset({str(c) for c in row})), None)
        row_map = {}
        if header_idx is not None:
            header = [str(c) for c in compliance_rows[header_idx]]
            indices = {name: header.index(name) for name in ["Supplier", *DIMENSIONS, "Compliance Score"]}
            for row in compliance_rows[header_idx + 1:]:
                if len(row) > indices["Supplier"] and row[indices["Supplier"]] is not None:
                    row_map[str(row[indices["Supplier"]]).strip()] = row
        compliance_ok = header_idx is not None and len(row_map) == len(records)
        for rec in records:
            row = row_map.get(rec["supplier"])
            if row is None:
                compliance_ok = False; continue
            if any(str(row[indices[dim]]).strip() != rec["statuses"][dim] for dim in DIMENSIONS): compliance_ok = False
            try:
                if abs(float(row[indices["Compliance Score"]]) - rec["score"]) > 0.01: compliance_ok = False
            except (TypeError, ValueError): compliance_ok = False
        results.append(criterion("R002", compliance_ok, f"重算记录={[(r['supplier'], r['score']) for r in records]}; 输出行={len(row_map)}"))

        action_rows = sheet_rows(args.artifact_dir / "corrective_actions.xlsx", "Corrective Actions")
        action_header = next((i for i, row in enumerate(action_rows) if {"Supplier", "Required Action", "Target Date"}.issubset({str(c) for c in row})), None)
        action_ok = action_header is not None
        good_actions = {rec["supplier"]: False for rec in records if rec["gaps"]}
        if action_header is not None:
            header = [str(c) for c in action_rows[action_header]]
            si, ai, di = header.index("Supplier"), header.index("Required Action"), header.index("Target Date")
            for row in action_rows[action_header + 1:]:
                if len(row) <= max(si, ai, di): continue
                supplier = str(row[si])
                if supplier not in good_actions: continue
                value = str(row[di])[:10]
                try: future = date.fromisoformat(value) >= date.today()
                except ValueError: future = False
                if str(row[ai]).strip() and future: good_actions[supplier] = True
        action_ok = action_ok and all(good_actions.values())
        results.append(criterion("R003", action_ok, f"整改覆盖={good_actions}"))

        findings = doc_text(args.artifact_dir / "audit_findings.docx")
        findings_low = findings.lower()
        findings_ok = all(rec["supplier"].lower() in findings_low for rec in records) and all((str(rec["score"]) in findings or f"{rec['score']:.2f}" in findings) for rec in records) and any(token in findings_low for token in ("sla", "liability", "ответствен", "confidentiality", "保密", "差距"))
        results.append(criterion("R004", findings_ok, "逐供应商报告和差距术语检查"))

        summary = doc_text(args.artifact_dir / "audit_summary.docx")
        summary_low = summary.lower()
        risk_counts = {risk: sum(rec["risk"] == risk for rec in records) for risk in ("High", "Medium", "Low")}
        avg = sum(rec["score"] for rec in records) / len(records)
        summary_ok = ("average" in summary_low or "средн" in summary_low) and f"{avg:.2f}" in summary and all(f"{risk}={count}" in summary for risk, count in risk_counts.items())
        results.append(criterion("R005", summary_ok, f"平均={avg:.2f}; 风险={risk_counts}"))

        email = (args.artifact_dir / "audit_notifications.eml").read_text(encoding="utf-8")
        emails_ok = all(rec["email"].lower() in email.lower() for rec in records) and "audit" in email.lower() and ("未发送" in email or "offline" in email.lower()) and all(rec["supplier"] in email for rec in records)
        results.append(criterion("R006", emails_ok, "收件人、主题和离线草稿标记检查"))

        events = parse_ics(args.artifact_dir / "supplier_review_meetings.ics")
        high_names = {rec["supplier"] for rec in records if rec["risk"] == "High"}
        event_text = "\n".join(event.get("SUMMARY", "") + " " + event.get("DESCRIPTION", "") for event in events)
        meetings_ok = len(events) == len(high_names) and all(name in event_text for name in high_names) and "Owner" in event_text and "Corrective" in event_text
        results.append(criterion("R007", meetings_ok, f"VEVENT={len(events)}; 高风险供应商={sorted(high_names)}"))

        cross_ok = all(rec["supplier"] in findings and f"{rec['score']:.2f}" in findings and rec["supplier"] in email for rec in records) and f"{avg:.2f}" in summary
        results.append(criterion("R008", cross_ok, "跨文件供应商/分数一致"))

        wb_native = all(zipfile.is_zipfile(args.artifact_dir / name) for name in ("compliance_assessment.xlsx", "corrective_actions.xlsx"))
        doc_native = all((args.artifact_dir / name).stat().st_size > 1000 for name in ("audit_findings.docx", "audit_summary.docx"))
        results.append(criterion("R009", wb_native and doc_native, f"OOXML={wb_native}; DOCX={doc_native}"))

        hashes_ok = all(digest(source_dir / filename) == expected for filename, expected in EXPECTED_HASHES.items())
        results.append(criterion("R010", hashes_ok, "输入副本 SHA256 检查"))
        reward = sum(item["score"] for item in results) / len(results)
    except Exception as exc:
        results.append(criterion("R001", False, f"Judge 读取失败：{exc}")); reward = 0.0
    passed = reward >= 0.70 and all(item["score"] == 1.0 for item in results[:5])
    payload = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "criterion_results": results, "reward": reward, "passed": passed}
    args.result.parent.mkdir(parents=True, exist_ok=True); args.result.write_text(__import__("json").dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.parent.mkdir(parents=True, exist_ok=True); args.reward_file.write_text(f"{reward:.6f}\n", encoding="utf-8")
    raise SystemExit(0 if passed else 1)


if __name__ == "__main__": main()
