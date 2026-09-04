#!/usr/bin/env python3
"""Reference construction for the offline supplier-contract audit workbook."""
from __future__ import annotations

import argparse
import csv
import hashlib
import re
import subprocess
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

CONTRACTS = [
    "contract_supplier1_TechService.pdf",
    "contract_supplier2_GlobalLogistics.pdf",
    "contract_supplier3_ControlQuality.pdf",
    "contract_supplier4_ConsultGroup.pdf",
    "contract_supplier5_ServicePodderzhka.pdf",
]
DIMENSIONS = ["Scope Clarity", "SLA Present", "Payment Terms", "IP Rights", "Confidentiality", "Liability Limits"]
HASHES = {
    "supplier_list.csv": "4553ce4cfd4ce8c336c900f433e1530d07c82b404a71eaa7e157ba9b71c78a1f",
    "contract_template.pdf": "f0b935f18a9a05772eb7b7f6bdaf98f7999ce9e55d50f4ee66452695e118df06",
    "contract_supplier1_TechService.pdf": "b54b1d5bc3a5e7bf459718e0a0aa2f337917a3646a1abb1f6ab725928f6edcc1",
    "contract_supplier2_GlobalLogistics.pdf": "b3f3fc8f20ed819c651dcb4435407680ea1da78785ccdaef02ab1d85cdddd98d",
    "contract_supplier3_ControlQuality.pdf": "066dfadfd2cf5a5e0d59e7e264fc1b0e0b4f5f7b8de50c38790c2646b1221fe8",
    "contract_supplier4_ConsultGroup.pdf": "9fbc24147b38bedca6202bf82ba2ae04df4ec098af0ccf6e6f2b142774ed8843",
    "contract_supplier5_ServicePodderzhka.pdf": "b7423c3e062654c76a1395685426e3a230be02fc7f2ede865dbc9ba515cd9907",
}


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def pdf_text(path: Path) -> str:
    result = subprocess.run(["pdftotext", "-layout", str(path), "-"], check=True, capture_output=True, text=True)
    return " ".join(result.stdout.lower().split())


def statuses(text: str) -> dict[str, str]:
    return {
        "Scope Clarity": "Incomplete" if "не детализированы" in text else "Complete",
        "SLA Present": "Complete" if all(token in text for token in ("99,5%", "4 часа", "24 часа")) else "Incomplete",
        "Payment Terms": "Compliant" if "в течение 30 дней" in text and "в течение 90 дней" not in text else "Non-Compliant",
        "IP Rights": "Compliant" if "все результаты работ" in text and "собственностью компании" in text else "Incomplete",
        "Confidentiality": "Compliant" if "конфиденциальность сведений" in text else "Missing",
        "Liability Limits": "Compliant" if "совокупная ответственность ограничена" in text else "Missing",
    }


def good(dimension: str, value: str) -> bool:
    return value == ("Complete" if dimension in {"Scope Clarity", "SLA Present"} else "Compliant")


def records(input_dir: Path) -> list[dict]:
    for name, expected in HASHES.items():
        if digest(input_dir / name) != expected:
            raise ValueError(f"input hash mismatch: {name}")
    with (input_dir / "supplier_list.csv").open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    if len(rows) != len(CONTRACTS):
        raise ValueError("supplier_list.csv must contain five rows")
    output = []
    for row, contract in zip(rows, CONTRACTS):
        state = statuses(pdf_text(input_dir / contract))
        gaps = [dim for dim in DIMENSIONS if not good(dim, state[dim])]
        score = round(100 * (len(DIMENSIONS) - len(gaps)) / len(DIMENSIONS), 2)
        risk = "High" if len(gaps) >= 2 else "Medium" if gaps else "Low"
        output.append({"supplier": row["Supplier_Name"], "contact": row["Primary_Contact"], "email": row["Contact_Email"], "service": row["Service_Type"], "start": row["Contract_Start_Date"], "value": float(row["Annual_Contract_Value"]), "contract": contract, "state": state, "gaps": gaps, "score": score, "risk": risk})
    return output


def style(ws, table_name: str, end_row: int, end_col: int) -> None:
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2, max_row=end_row, max_col=end_col):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if end_row >= 2:
        table = Table(displayName=table_name, ref=f"A1:{chr(64 + end_col)}{end_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        ws.add_table(table)
    for column in ws.columns:
        letter = column[0].column_letter
        ws.column_dimensions[letter].width = min(48, max(14, max(len(str(cell.value or "")) for cell in column) + 2))


def build(input_dir: Path, output_dir: Path) -> None:
    recs = records(input_dir)
    wb = Workbook()
    register = wb.active
    register.title = "Supplier Register"
    register.append(["Supplier", "Primary Contact", "Contact Email", "Service Type", "Contract Start Date", "Annual Contract Value", "Contract File"])
    for rec in recs:
        register.append([rec["supplier"], rec["contact"], rec["email"], rec["service"], rec["start"], rec["value"], rec["contract"]])
    style(register, "SupplierRegister", register.max_row, 7)
    for row in range(2, register.max_row + 1): register.cell(row, 6).number_format = '#,##0.00'

    assessment = wb.create_sheet("Compliance Assessment")
    assessment.append(["Supplier", *DIMENSIONS, "Compliance Score", "Risk"])
    for idx, rec in enumerate(recs, 2):
        assessment.append([rec["supplier"], *[rec["state"][dim] for dim in DIMENSIONS], f'=ROUND(COUNTIF(B{idx}:G{idx},"Complete")+COUNTIF(B{idx}:G{idx},"Compliant"),0)/6*100', f'=IF(COUNTIF(B{idx}:G{idx},"Incomplete")+COUNTIF(B{idx}:G{idx},"Non-Compliant")+COUNTIF(B{idx}:G{idx},"Missing")=0,"Low",IF(COUNTIF(B{idx}:G{idx},"Incomplete")+COUNTIF(B{idx}:G{idx},"Non-Compliant")+COUNTIF(B{idx}:G{idx},"Missing")>=2,"High","Medium"))'])
    style(assessment, "ComplianceAssessment", assessment.max_row, 9)
    for row in range(2, assessment.max_row + 1): assessment.cell(row, 8).number_format = "0.00"

    actions = wb.create_sheet("Corrective Actions")
    actions.append(["Supplier", "Issue", "Required Action", "Owner", "Target Date", "Est. Cost", "Status"])
    action_text = {"Scope Clarity": ("Scope Clarity gap", "补充服务范围、交付成果和期限"), "SLA Present": ("SLA Present gap", "补充可衡量的可用性、响应和修复指标"), "Payment Terms": ("Payment Terms gap", "复核并修订付款条款"), "IP Rights": ("IP Rights gap", "明确公司对合同成果的权利"), "Confidentiality": ("Confidentiality gap", "补充保密义务和期限"), "Liability Limits": ("Liability Limits gap", "补充责任上限条款")}
    for ridx, rec in enumerate(recs):
        for gidx, dim in enumerate(rec["gaps"]):
            issue, action = action_text[dim]
            actions.append([rec["supplier"], issue, action, "Procurement Compliance", (date.today() + timedelta(days=21 + ridx * 7 + gidx * 7)).isoformat(), 0, "Planned"])
    style(actions, "CorrectiveActions", actions.max_row, 7)
    for row in range(2, actions.max_row + 1): actions.cell(row, 6).number_format = '#,##0.00'

    findings = wb.create_sheet("Audit Findings")
    findings.append(["Supplier", "Compliance Score", "Risk", "Findings", "Evidence", "Recommendation"])
    for rec in recs:
        gaps = "; ".join(f"{dim}={rec['state'][dim]}" for dim in rec["gaps"]) or "No gaps found across the six template dimensions"
        evidence = "; ".join(rec["gaps"]) if rec["gaps"] else "All six template sections present"
        recommendation = "; ".join(action_text[d][1] for d in rec["gaps"]) or "Continue monitoring against the template"
        findings.append([rec["supplier"], f"='Compliance Assessment'!H{findings.max_row + 1}", f"='Compliance Assessment'!I{findings.max_row + 1}", gaps, evidence, recommendation])
    style(findings, "AuditFindings", findings.max_row, 6)
    for row in range(2, findings.max_row + 1): findings.cell(row, 2).number_format = "0.00"

    summary = wb.create_sheet("Audit Summary")
    summary.append(["Metric", "Value"])
    summary.append(["Average Compliance Score", "=AVERAGE('Compliance Assessment'!H2:H6)"])
    summary.append(["High Risk Suppliers", '=COUNTIF(\'Compliance Assessment\'!I2:I6,"High")'])
    summary.append(["Medium Risk Suppliers", '=COUNTIF(\'Compliance Assessment\'!I2:I6,"Medium")'])
    summary.append(["Low Risk Suppliers", '=COUNTIF(\'Compliance Assessment\'!I2:I6,"Low")'])
    summary.append(["High Risk Supplier Names", '=IF(\'Compliance Assessment\'!I2="High",\'Compliance Assessment\'!A2&", ","")&IF(\'Compliance Assessment\'!I3="High",\'Compliance Assessment\'!A3&", ","")&IF(\'Compliance Assessment\'!I4="High",\'Compliance Assessment\'!A4&", ","")&IF(\'Compliance Assessment\'!I5="High",\'Compliance Assessment\'!A5&", ","")&IF(\'Compliance Assessment\'!I6="High",\'Compliance Assessment\'!A6,"" )'])
    style(summary, "AuditSummary", summary.max_row, 2)
    summary.cell(2, 2).number_format = "0.00"

    notices = wb.create_sheet("Notification Drafts")
    notices.append(["Supplier", "To", "Subject", "Draft Body", "Status"])
    for rec in recs:
        gaps = ", ".join(rec["gaps"]) if rec["gaps"] else "No corrective action required"
        notices.append([rec["supplier"], f"='Supplier Register'!C{notices.max_row + 1}", "Contract Compliance Audit Findings", f"Compliance Score {rec['score']:.2f}/100; Risk {rec['risk']}; Corrective actions: {gaps}", "Draft - Not Sent"])
    style(notices, "NotificationDrafts", notices.max_row, 5)

    meetings = wb.create_sheet("Review Meetings")
    meetings.append(["Supplier", "Meeting Title", "Owner", "Corrective Actions", "Target Date", "Status"])
    for rec in recs:
        if rec["risk"] != "High": continue
        gaps = ", ".join(rec["gaps"])
        meetings.append([rec["supplier"], "Contract Compliance Improvement Review", "Procurement Compliance", gaps, (date.today() + timedelta(days=30)).isoformat(), "Draft - Not Created"])
    style(meetings, "ReviewMeetings", meetings.max_row, 6)

    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    output_dir.mkdir(parents=True, exist_ok=True)
    wb.save(output_dir / "supplier_contract_audit.xlsx")
    wb.close()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    build(args.input_dir, args.output_dir)


if __name__ == "__main__":
    main()
