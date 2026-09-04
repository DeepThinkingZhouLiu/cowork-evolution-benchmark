#!/usr/bin/env python3
"""Independent deterministic Judge for the XLSX-only contract audit."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

OUTPUT = "supplier_contract_audit.xlsx"
SHEETS = ["Supplier Register", "Compliance Assessment", "Corrective Actions", "Audit Findings", "Audit Summary", "Notification Drafts", "Review Meetings"]
CONTRACTS = ["contract_supplier1_TechService.pdf", "contract_supplier2_GlobalLogistics.pdf", "contract_supplier3_ControlQuality.pdf", "contract_supplier4_ConsultGroup.pdf", "contract_supplier5_ServicePodderzhka.pdf"]
DIMENSIONS = ["Scope Clarity", "SLA Present", "Payment Terms", "IP Rights", "Confidentiality", "Liability Limits"]
HASHES = {
    "supplier_list.csv": "4553ce4cfd4ce8c336c900f433e1530d07c82b404a71eaa7e157ba9b71c78a1f",
    "contract_template.pdf": "f0b935f18a9a05772eb7b7f6bdaf98f7999ce9e55d50f4ee66452695e118df06",
    "contract_supplier1_TechService.pdf": "b54b1d5bc3a5e7bf459718e0a0aa2f337917a3646a1abb1f6ab725928f6edcc1",
    "contract_supplier2_GlobalLogistics.pdf": "b3f3fc8f20ed819c651dcb4435407680ea1da78785ccdaef02ab1d85cdddd98d",
    "contract_supplier3_ControlQuality.pdf": "066dfadfd2cf5a5e0d59e7e264fc1b0e0b4f5f7b8de50c38790c2646b1221fe8",
    "contract_supplier4_ConsultGroup.pdf": "9fbc24147b38bedca6202bf82ba2ae04df4ec098af0ccf6e6f2b142774ed8843",
    "contract_supplier5_ServicePodderzhka.pdf": "b7423c3e062654c76a1395685426e3a230be02fc7f2ede865dbc9ba515cd9907",
}
WEIGHTS = {"R001": 10, "R002": 15, "R003": 25, "R004": 15, "R005": 10, "R006": 10, "R007": 5, "R008": 5, "R009": 5, "R010": 5, "R011": -20}
KINDS = {"R001": "hurdle", "R002": "hurdle", "R003": "hurdle", "R004": "hurdle", "R005": "hurdle", "R006": "hurdle", "R007": "positive", "R008": "positive", "R009": "positive", "R010": "positive", "R011": "penalty"}
CAPS = {"R001": 0.0, "R002": 0.5, "R003": 0.45, "R004": 0.6, "R005": 0.65, "R006": 0.7}


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def pdf_text(path: Path) -> str:
    raw = subprocess.run(["pdftotext", "-layout", str(path), "-"], check=True, capture_output=True, text=True).stdout.lower()
    return " ".join(raw.split())


def expected_records(source: Path) -> list[dict]:
    with (source / "supplier_list.csv").open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    records = []
    for row, filename in zip(rows, CONTRACTS):
        text = pdf_text(source / filename)
        state = {
            "Scope Clarity": "Incomplete" if "не детализированы" in text else "Complete",
            "SLA Present": "Complete" if all(x in text for x in ("99,5%", "4 часа", "24 часа")) else "Incomplete",
            "Payment Terms": "Compliant" if "в течение 30 дней" in text and "в течение 90 дней" not in text else "Non-Compliant",
            "IP Rights": "Compliant" if "все результаты работ" in text and "собственностью компании" in text else "Incomplete",
            "Confidentiality": "Compliant" if "конфиденциальность сведений" in text else "Missing",
            "Liability Limits": "Compliant" if "совокупная ответственность ограничена" in text else "Missing",
        }
        good = lambda dim: state[dim] == ("Complete" if dim in {"Scope Clarity", "SLA Present"} else "Compliant")
        gaps = [dim for dim in DIMENSIONS if not good(dim)]
        score = round(100 * (len(DIMENSIONS) - len(gaps)) / 6, 2)
        records.append({"supplier": row["Supplier_Name"], "contact": row["Primary_Contact"], "email": row["Contact_Email"], "service": row["Service_Type"], "start": row["Contract_Start_Date"], "value": float(row["Annual_Contract_Value"]), "contract": filename, "state": state, "gaps": gaps, "score": score, "risk": "High" if len(gaps) >= 2 else "Medium" if gaps else "Low"})
    return records


def recalc(path: Path) -> Path | None:
    binary = shutil.which("libreoffice") or shutil.which("soffice")
    if not binary:
        return None
    directory = Path(tempfile.mkdtemp(prefix="supplier-audit-recalc-"))
    source_dir = directory / "source"
    output_dir = directory / "output"
    source_dir.mkdir(); output_dir.mkdir()
    copied = source_dir / path.name
    shutil.copy2(path, copied)
    profile = directory / "profile"
    run = subprocess.run([binary, "--headless", f"-env:UserInstallation={profile.resolve().as_uri()}", "--convert-to", "xlsx", "--outdir", str(output_dir), str(copied)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=240, check=False)
    result = output_dir / path.name
    return result if run.returncode == 0 and result.is_file() else None


def criterion(cid: str, score: float, evidence: str) -> dict:
    return {"criterion_id": cid, "score": float(score), "evidence": evidence, "method": "deterministic", "weight": WEIGHTS[cid], "type": KINDS[cid]}


def rows(wb, name: str) -> list[tuple]:
    return list(wb[name].iter_rows(values_only=True))


def header_index(table: list[tuple], required: set[str]) -> tuple[int, dict[str, int]] | None:
    for idx, row in enumerate(table):
        values = {str(v): i for i, v in enumerate(row) if v is not None}
        if required.issubset(values):
            return idx, values
    return None


def evaluate(artifact_dir: Path, render: bool = True) -> dict:
    source = Path(__file__).parents[1] / "data/input_files"
    output = artifact_dir / OUTPUT
    expected = expected_records(source)
    result: list[dict] = []
    names = sorted(p.name for p in artifact_dir.iterdir()) if artifact_dir.is_dir() else []
    valid = output.is_file() and output.stat().st_size > 0 and zipfile.is_zipfile(output)
    workbook = None
    cached = None
    if valid:
        try:
            workbook = load_workbook(output, data_only=False)
            valid = workbook.sheetnames == SHEETS
        except Exception:
            valid = False
    result.append(criterion("R001", 1 if valid and names == [OUTPUT] else 0, f"valid={valid}; files={names}; sheets={workbook.sheetnames if workbook else None}"))
    if not valid:
        for cid in ("R002", "R003", "R004", "R005", "R006"):
            result.append(criterion(cid, 0, "invalid workbook"))
        result.extend([criterion("R007", 0, "invalid workbook"), criterion("R008", 0, "invalid workbook"), criterion("R009", 0, "invalid workbook"), criterion("R010", 1 if all(digest(source / n) == h for n, h in HASHES.items()) else 0, "input hashes"), criterion("R011", 1, "invalid output")])
        return finalize(result, output)

    recalculated = recalc(output)
    if recalculated:
        cached = load_workbook(recalculated, data_only=True)

    register = rows(workbook, "Supplier Register")
    reg_head = header_index(register, {"Supplier", "Primary Contact", "Contact Email", "Service Type", "Contract Start Date", "Annual Contract Value", "Contract File"})
    reg_ok = bool(reg_head)
    if reg_head:
        hidx, hi = reg_head
        data = register[hidx + 1:hidx + 1 + len(expected)]
        reg_ok = len(data) == len(expected)
        for rec, row in zip(expected, data):
            checks = [row[hi["Supplier"]] == rec["supplier"], row[hi["Primary Contact"]] == rec["contact"], row[hi["Contact Email"]] == rec["email"], row[hi["Service Type"]] == rec["service"], str(row[hi["Contract Start Date"]])[:10] == rec["start"], abs(float(row[hi["Annual Contract Value"]]) - rec["value"]) < 0.01, row[hi["Contract File"]] == rec["contract"]]
            reg_ok &= all(checks)
    result.append(criterion("R002", 1 if reg_ok else 0, f"rows={len(register)}; exact CSV fields={reg_ok}"))

    assessment = rows(workbook, "Compliance Assessment")
    ah = header_index(assessment, {"Supplier", *DIMENSIONS, "Compliance Score", "Risk"})
    assessment_ok = bool(ah)
    if ah:
        aidx, ai = ah; data = assessment[aidx + 1:aidx + 1 + len(expected)]
        assessment_ok = len(data) == len(expected)
        for rec, row in zip(expected, data):
            assessment_ok &= row[ai["Supplier"]] == rec["supplier"]
            assessment_ok &= all(str(row[ai[d]]).strip() == rec["state"][d] for d in DIMENSIONS)
            score_value = (cached["Compliance Assessment"].cell(aidx + 2 + expected.index(rec), ai["Compliance Score"] + 1).value if cached else row[ai["Compliance Score"]])
            risk_value = (cached["Compliance Assessment"].cell(aidx + 2 + expected.index(rec), ai["Risk"] + 1).value if cached else row[ai["Risk"]])
            assessment_ok &= score_value is not None and abs(float(score_value) - rec["score"]) <= 0.02 and str(risk_value) == rec["risk"]
    result.append(criterion("R003", 1 if assessment_ok else 0, f"independently recomputed statuses/scores={assessment_ok}"))

    actions = rows(workbook, "Corrective Actions")
    ac = header_index(actions, {"Supplier", "Issue", "Required Action", "Target Date"})
    action_flags = {rec["supplier"]: False for rec in expected if rec["gaps"]}
    if ac:
        idx, ai = ac
        for row in actions[idx + 1:]:
            if not row or row[ai["Supplier"]] is None: continue
            supplier = str(row[ai["Supplier"]]); action = str(row[ai["Required Action"]] or "").strip(); target = str(row[ai["Target Date"]] or "")[:10]
            if supplier not in action_flags or not action: continue
            try: future = date.fromisoformat(target) >= date.today()
            except ValueError: future = False
            if future: action_flags[supplier] = True
    action_ok = bool(ac) and all(action_flags.values())
    result.append(criterion("R004", 1 if action_ok else 0, f"coverage={action_flags}"))

    findings = rows(workbook, "Audit Findings")
    fh = header_index(findings, {"Supplier", "Compliance Score", "Risk", "Findings", "Evidence", "Recommendation"})
    findings_ok = bool(fh)
    if fh:
        idx, fi = fh; data = findings[idx + 1:]
        by_supplier = {str(row[fi["Supplier"]]): row for row in data if row[fi["Supplier"]] is not None}
        findings_ok = len(by_supplier) == len(expected)
        for rec in expected:
            row = by_supplier.get(rec["supplier"]); text = " ".join(str(row[fi[k]] or "") for k in ("Findings", "Evidence", "Recommendation")) if row else ""
            score_cell = (cached["Audit Findings"].cell(idx + 2 + expected.index(rec), fi["Compliance Score"] + 1).value if cached else row[fi["Compliance Score"]]) if row is not None else None
            score_matches = f"{rec['score']:.2f}" in text
            if not score_matches:
                try: score_matches = abs(float(score_cell) - rec["score"]) < .02
                except (TypeError, ValueError): score_matches = False
            findings_ok &= row is not None and score_matches
            risk_cell = (cached["Audit Findings"].cell(idx + 2 + expected.index(rec), fi["Risk"] + 1).value if cached else row[fi["Risk"]]) if row is not None else None
            findings_ok &= row is not None and str(risk_cell) == rec["risk"] and (not rec["gaps"] or any(dim in text for dim in rec["gaps"]))
    result.append(criterion("R005", 1 if findings_ok else 0, f"five supplier findings={findings_ok}"))

    summary = rows(workbook, "Audit Summary")
    sh = header_index(summary, {"Metric", "Value"})
    summary_ok = bool(sh)
    if sh:
        idx, si = sh; mapping = {str(row[si["Metric"]]): row[si["Value"]] for row in summary[idx + 1:] if row[si["Metric"]] is not None}
        avg = sum(rec["score"] for rec in expected) / len(expected)
        summary_ok = "Average Compliance Score" in mapping and "High Risk Suppliers" in mapping and "Medium Risk Suppliers" in mapping and "Low Risk Suppliers" in mapping
        if cached:
            cached_summary = cached["Audit Summary"]
            for row_number, expected_value in ((2, avg), (3, 2), (4, 2), (5, 1)):
                value = cached_summary.cell(row_number, 2).value
                summary_ok &= value is not None and abs(float(value) - expected_value) <= .02
        summary_ok &= all(isinstance(mapping[k], str) and str(mapping[k]).startswith("=") for k in ("Average Compliance Score", "High Risk Suppliers", "Medium Risk Suppliers", "Low Risk Suppliers"))
    result.append(criterion("R006", 1 if summary_ok else 0, f"formula summary and independent aggregation={summary_ok}"))

    notices = rows(workbook, "Notification Drafts")
    nh = header_index(notices, {"Supplier", "To", "Subject", "Draft Body", "Status"})
    notices_ok = bool(nh)
    if nh:
        idx, ni = nh; data = notices[idx + 1:idx + 1 + len(expected)]
        notices_ok = len(data) == len(expected)
        for rec, row in zip(expected, data):
            body = str(row[ni["Draft Body"]] or "")
            to = (cached["Notification Drafts"].cell(idx + 2 + expected.index(rec), ni["To"] + 1).value if cached else row[ni["To"]])
            notices_ok &= rec["supplier"] in str(row[ni["Supplier"]]) and rec["email"].lower() == str(to).lower() and "Contract Compliance Audit" in str(row[ni["Subject"]]) and f"{rec['score']:.2f}" in body and rec["risk"] in body and str(row[ni["Status"]]) == "Draft - Not Sent"
    result.append(criterion("R007", 1 if notices_ok else 0, f"five offline notices={notices_ok}"))

    meetings = rows(workbook, "Review Meetings")
    mh = header_index(meetings, {"Supplier", "Meeting Title", "Owner", "Corrective Actions", "Target Date", "Status"})
    high = [rec for rec in expected if rec["risk"] == "High"]
    meetings_ok = bool(mh)
    if mh:
        idx, mi = mh; data = meetings[idx + 1:]
        by_supplier = {str(row[mi["Supplier"]]): row for row in data if row[mi["Supplier"]] is not None}
        meetings_ok = set(by_supplier) == {rec["supplier"] for rec in high}
        for rec in high:
            row = by_supplier.get(rec["supplier"]); meetings_ok &= row is not None and "Compliance" in str(row[mi["Meeting Title"]]) and str(row[mi["Owner"]]).strip() and all(dim in str(row[mi["Corrective Actions"]]) for dim in rec["gaps"]) and str(row[mi["Status"]]) == "Draft - Not Created"
    result.append(criterion("R008", 1 if meetings_ok else 0, f"high-risk meeting drafts={meetings_ok}"))

    calc = workbook.calculation
    formula_ok = all(isinstance(workbook["Compliance Assessment"].cell(row, col).value, str) and workbook["Compliance Assessment"].cell(row, col).value.startswith("=") for row in range(2, 7) for col in (8, 9)) and all(isinstance(workbook["Audit Summary"].cell(row, 2).value, str) and workbook["Audit Summary"].cell(row, 2).value.startswith("=") for row in range(2, 6))
    structure_ok = workbook.sheetnames == SHEETS and workbook["Compliance Assessment"].freeze_panes == "A2" and workbook["Audit Summary"].freeze_panes == "A2" and getattr(calc, "calcMode", None) == "auto" and bool(getattr(calc, "fullCalcOnLoad", False)) and bool(getattr(calc, "forceFullCalc", False))
    hashes_ok = all(digest(source / name) == expected_hash for name, expected_hash in HASHES.items())
    result.append(criterion("R010", 1 if hashes_ok else 0, f"input_sha256={hashes_ok}"))
    unchanged_input = digest(output) == digest(source / "supplier_list.csv")
    result.append(criterion("R011", 1 if unchanged_input else 0, "unchanged input penalty triggered" if unchanged_input else "output is not an unchanged input"))
    if render:
        binary = shutil.which("libreoffice") or shutil.which("soffice")
        pdfinfo = shutil.which("pdfinfo")
        render_ok = False
        evidence = "missing LibreOffice/pdfinfo"
        if binary and pdfinfo:
            directory = Path(tempfile.mkdtemp(prefix="supplier-audit-render-")); profile = directory / "profile"
            run = subprocess.run([binary, "--headless", f"-env:UserInstallation={profile.resolve().as_uri()}", "--convert-to", "pdf", "--outdir", str(directory), str(output)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=240, check=False)
            pdf = directory / "supplier_contract_audit.pdf"
            if run.returncode == 0 and pdf.is_file() and pdf.stat().st_size > 1000:
                info = subprocess.run([pdfinfo, str(pdf)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, check=False).stdout
                match = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
                render_ok = bool(match and int(match.group(1)) >= len(SHEETS))
                evidence = f"PDF={pdf.stat().st_size}; {match.group(0) if match else 'no pages'}"
        render_ok = render_ok
    else:
        render_ok, evidence = True, "render skipped"
    result.append(criterion("R009", 1 if formula_ok and structure_ok and render_ok else 0, f"native formulas={formula_ok}; structure/recalc={structure_ok}; {evidence}"))
    if cached: cached.close()
    workbook.close()
    return finalize(result, output)


def finalize(results: list[dict], output: Path) -> dict:
    positive_total = sum(weight for weight in WEIGHTS.values() if weight > 0)
    earned = sum(item["score"] * WEIGHTS[item["criterion_id"]] for item in results if WEIGHTS[item["criterion_id"]] > 0)
    penalties = sum(WEIGHTS[item["criterion_id"]] for item in results if WEIGHTS[item["criterion_id"]] < 0 and item["score"] >= 1)
    reward = max(0.0, min(1.0, (earned + penalties) / positive_total))
    failures = [item["criterion_id"] for item in results if KINDS[item["criterion_id"]] == "hurdle" and item["score"] < 1]
    if failures: reward = min(reward, min(CAPS[cid] for cid in failures))
    return {"task_id": "cowork-evo/cowork-bench-supplier-contract-workbook", "rubric_version": "cowork-bench-supplier-contract-workbook-v0.1", "judge_version": "2.0.0", "artifact_sha256": [digest(output)] if output.is_file() else [], "criterion_results": [{k: item[k] for k in ("criterion_id", "score", "evidence", "method")} for item in sorted(results, key=lambda x: x["criterion_id"])], "reward": round(reward, 6), "passed": reward >= .8 and not failures, "hurdle_failures": failures}


def main() -> None:
    parser = argparse.ArgumentParser(); parser.add_argument("--artifact-dir", type=Path, required=True); parser.add_argument("--result", type=Path, required=True); parser.add_argument("--reward-file", type=Path, required=True); parser.add_argument("--skip-render", action="store_true")
    args = parser.parse_args()
    payload = evaluate(args.artifact_dir.resolve(), not args.skip_render)
    args.result.parent.mkdir(parents=True, exist_ok=True); args.result.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.parent.mkdir(parents=True, exist_ok=True); args.reward_file.write_text(f"{payload['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    raise SystemExit(0 if payload["passed"] else 1)


if __name__ == "__main__":
    main()
