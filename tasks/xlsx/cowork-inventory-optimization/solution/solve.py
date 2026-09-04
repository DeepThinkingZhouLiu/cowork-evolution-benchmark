#!/usr/bin/env python3
"""Build the offline inventory-optimization workbook from the frozen inputs."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


OUTPUT_NAME = "Inventory_Optimization_Report.xlsx"
INPUT_HASHES = {
    "demand_history.csv": "9c8cfde544c5fbb0a8bdfb8112424fa92b4b0581077ea6d83535034a8e1397c7",
    "snapshot_date.txt": "6dbe997331c8846eb24631dc39833916a796f8cf36236d558e072f185aa3d4e8",
    "store_orders.json": "e34e4f591a4dc04620f8986d2d339f1865c32d1f0fd4ab3c498be6a1b294baa3",
    "store_products.json": "4ef055016ee64285773250bbafb7413ab2c04c6c19a93b2a9d6d31fd2b3124aa",
    "supplier_catalog.json": "bb7b9b3710e86299de06fc116d5b839809fad563e9ba5316be59e7bdc5a012f9",
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_inputs(input_dir: Path) -> tuple[list[dict], list[dict], dict, list[dict], str]:
    for name, expected in INPUT_HASHES.items():
        actual = sha256(input_dir / name)
        if actual != expected:
            raise ValueError(f"input hash mismatch for {name}: {actual}")
    products = json.loads((input_dir / "store_products.json").read_text(encoding="utf-8"))
    orders = json.loads((input_dir / "store_orders.json").read_text(encoding="utf-8"))
    catalog = json.loads((input_dir / "supplier_catalog.json").read_text(encoding="utf-8"))
    with (input_dir / "demand_history.csv").open(encoding="utf-8-sig", newline="") as stream:
        demand = list(csv.DictReader(stream))
    snapshot = (input_dir / "snapshot_date.txt").read_text(encoding="utf-8").strip()
    return products, orders, catalog, demand, snapshot


def style_sheet(sheet, table_name: str, max_col: int) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=max_col):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if sheet.max_row >= 2:
        table = Table(displayName=table_name, ref=f"A1:{get_column_letter(max_col)}{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)
    for column in sheet.columns:
        letter = column[0].column_letter
        longest = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[letter].width = min(54, max(14, longest + 2))


def build(input_dir: Path, output_dir: Path) -> Path:
    products, orders, catalog, demand, snapshot = load_inputs(input_dir)
    eligible = [
        product
        for product in products
        if product.get("status") == "publish" and int(product.get("current_stock", 0)) > 0
    ]
    sold: defaultdict[int, int] = defaultdict(int)
    for order in orders:
        if order.get("status") not in {"completed", "processing"}:
            continue
        for item in order.get("line_items", []):
            sold[int(item["product_id"])] += int(item["quantity"])
    eligible.sort(key=lambda product: (str(product.get("product_name", "")).casefold(), int(product["product_id"])))

    workbook = Workbook()
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    stock = workbook.active
    stock.title = "Stock_Status"
    stock.append(["Product_ID", "Product", "Category", "Current_Stock", "Total_Sales", "Daily_Rate", "Days_Remaining", "Needs_Restock"])
    stock_row_by_id: dict[int, int] = {}
    for product in eligible:
        product_id = int(product["product_id"])
        stock.append([
            product_id,
            product["product_name"],
            product["categories"][0]["name"],
            int(product["current_stock"]),
            sold[product_id],
            f"=E{stock.max_row + 1}/90",
            f"=IF(F{stock.max_row + 1}=0,999,ROUND(D{stock.max_row + 1}/F{stock.max_row + 1},1))",
            f'=IF(G{stock.max_row + 1}<30,"Yes","No")',
        ])
        stock_row_by_id[product_id] = stock.max_row
    style_sheet(stock, "StockStatus", 8)
    for row in range(2, stock.max_row + 1):
        stock.cell(row, 6).number_format = "0.00"
        stock.cell(row, 7).number_format = "0.0"

    supplier = workbook.create_sheet("Supplier_Info")
    supplier.append(["Supplier", "City", "Lead_Time_Days", "Min_Order_Qty", "Reliability_Score", "Risk_Index", "Risk_Band"])
    for item in catalog.get("suppliers", []):
        row = supplier.max_row + 1
        supplier.append([
            item["name"],
            item.get("city", ""),
            int(item["lead_time_days"]),
            int(item["min_order_qty"]),
            float(item["reliability_score"]),
            f"=ROUND(C{row}*(5-E{row}),2)",
            f'=IF(OR(C{row}>=14,E{row}<4),"High",IF(OR(C{row}>=10,E{row}<4.5),"Medium","Low"))',
        ])
    style_sheet(supplier, "SupplierInfo", 7)
    for row in range(2, supplier.max_row + 1):
        supplier.cell(row, 5).number_format = "0.0"
        supplier.cell(row, 6).number_format = "0.00"

    history = workbook.create_sheet("Demand_History")
    history.append(["Category", "Total_Units_Ordered", "Order_Count"])
    for row in demand:
        history.append([row["Category"], int(row["Total_Units_Ordered"]), int(row["Order_Count"])])
    style_sheet(history, "DemandHistory", 3)

    plan = workbook.create_sheet("Restock_Plan")
    plan.append(["Priority_Rank", "Product_ID", "Product", "Category", "Current_Stock", "Total_Sales", "Days_Remaining", "Recommended_Action"])
    # The plan is a view over Stock_Status, so every product fact remains linked.
    restock_ids = []
    for product in eligible:
        product_id = int(product["product_id"])
        daily = sold[product_id] / 90 if sold[product_id] else 0
        days = product["current_stock"] / daily if daily else 999
        if days < 30:
            restock_ids.append((days, product_id))
    restock_ids.sort(key=lambda pair: (pair[0], pair[1]))
    for rank, (_, product_id) in enumerate(restock_ids, 1):
        source_row = stock_row_by_id[product_id]
        plan.append([
            f"=ROW()-1",
            f"='Stock_Status'!A{source_row}",
            f"='Stock_Status'!B{source_row}",
            f"='Stock_Status'!C{source_row}",
            f"='Stock_Status'!D{source_row}",
            f"='Stock_Status'!E{source_row}",
            f"='Stock_Status'!G{source_row}",
            "Place replenishment order; review supplier lead time and reliability",
        ])
    style_sheet(plan, "RestockPlan", 8)
    for row in range(2, plan.max_row + 1):
        plan.cell(row, 7).number_format = "0.0"

    summary = workbook.create_sheet("Restock_Summary")
    summary.append(["Metric", "Value"])
    last_stock = stock.max_row
    summary.append(["Total_Products_Analyzed", f"=COUNTA('Stock_Status'!A2:A{last_stock})"])
    summary.append(["Products_Need_Restock", f'=COUNTIF(\'Stock_Status\'!H2:H{last_stock},"Yes")'])
    summary.append(["Products_Healthy", f'=COUNTIF(\'Stock_Status\'!H2:H{last_stock},"No")'])
    summary.append(["Avg_Days_Remaining", f"=ROUND(AVERAGE('Stock_Status'!G2:G{last_stock}),1)"])
    summary.append(["Snapshot_Date", snapshot])
    style_sheet(summary, "RestockSummary", 2)
    summary.cell(5, 2).number_format = "0.0"

    output_dir.mkdir(parents=True, exist_ok=True)
    output = output_dir / OUTPUT_NAME
    workbook.save(output)
    workbook.close()
    return output


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=Path("/app/data/input_files"))
    parser.add_argument("--output-dir", type=Path, default=Path("/app/output"))
    args = parser.parse_args()
    build(args.input_dir, args.output_dir)


if __name__ == "__main__":
    main()
