#!/usr/bin/env python3
"""Independent deterministic Judge for the offline inventory workbook."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import shutil
import subprocess
import tempfile
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


OUTPUT = "Inventory_Optimization_Report.xlsx"
SHEETS = ["Stock_Status", "Supplier_Info", "Demand_History", "Restock_Plan", "Restock_Summary"]
INPUT_HASHES = {
    "demand_history.csv": "9c8cfde544c5fbb0a8bdfb8112424fa92b4b0581077ea6d83535034a8e1397c7",
    "snapshot_date.txt": "6dbe997331c8846eb24631dc39833916a796f8cf36236d558e072f185aa3d4e8",
    "store_orders.json": "e34e4f591a4dc04620f8986d2d339f1865c32d1f0fd4ab3c498be6a1b294baa3",
    "store_products.json": "4ef055016ee64285773250bbafb7413ab2c04c6c19a93b2a9d6d31fd2b3124aa",
    "supplier_catalog.json": "bb7b9b3710e86299de06fc116d5b839809fad563e9ba5316be59e7bdc5a012f9",
}
WEIGHTS = {"R001": 10, "R002": 20, "R003": 20, "R004": 10, "R005": 10, "R006": 10, "R007": 10, "R008": 5, "R009": 5, "R010": 5, "R011": -20}
KINDS = {"R001": "hurdle", "R002": "hurdle", "R003": "hurdle", "R004": "hurdle", "R005": "hurdle", "R006": "hurdle", "R007": "hurdle", "R008": "positive", "R009": "positive", "R010": "positive", "R011": "penalty"}
CAPS = {"R001": 0.0, "R002": 0.55, "R003": 0.6, "R004": 0.65, "R005": 0.65, "R006": 0.65, "R007": 0.7}


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def criterion(cid: str, passed: bool, evidence: str = "") -> dict:
    return {"criterion_id": cid, "score": 1.0 if passed else 0.0, "evidence": evidence, "method": "deterministic"}


def normalize(value) -> str:
    return " ".join(str(value or "").split()).strip().casefold()


def number(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def inputs(source: Path) -> tuple[list[dict], list[dict], list[dict], list[dict], str]:
    products = json.loads((source / "store_products.json").read_text(encoding="utf-8"))
    orders = json.loads((source / "store_orders.json").read_text(encoding="utf-8"))
    catalog = json.loads((source / "supplier_catalog.json").read_text(encoding="utf-8"))["suppliers"]
    with (source / "demand_history.csv").open(encoding="utf-8-sig", newline="") as stream:
        demand = list(csv.DictReader(stream))
    snapshot = (source / "snapshot_date.txt").read_text(encoding="utf-8").strip()
    return products, orders, catalog, demand, snapshot


def expected_records(source: Path):
    products, orders, catalog, demand_csv, snapshot = inputs(source)
    category_by_id = {
        int(product["product_id"]): product.get("categories", [{}])[0].get("name", "")
        for product in products
    }
    sold: defaultdict[int, int] = defaultdict(int)
    demand_units: defaultdict[str, int] = defaultdict(int)
    demand_lines: defaultdict[str, int] = defaultdict(int)
    for order in orders:
        if order.get("status") not in {"completed", "processing"}:
            continue
        for item in order.get("line_items", []):
            product_id = int(item["product_id"])
            quantity = int(item["quantity"])
            sold[product_id] += quantity
            category = category_by_id.get(product_id, "")
            demand_units[category] += quantity
            demand_lines[category] += 1
    eligible = [
        product for product in products
        if product.get("status") == "publish" and int(product.get("current_stock", 0)) > 0
    ]
    eligible.sort(key=lambda product: (str(product.get("product_name", "")).casefold(), int(product["product_id"])))
    stock = []
    for product in eligible:
        product_id = int(product["product_id"])
        total = sold[product_id]
        exact_daily = total / 90
        daily = round(exact_daily, 2)
        days = 999 if exact_daily == 0 else round(int(product["current_stock"]) / exact_daily, 1)
        stock.append({
            "id": product_id,
            "name": product["product_name"],
            "category": product.get("categories", [{}])[0].get("name", ""),
            "stock": int(product["current_stock"]),
            "sales": total,
            "daily": daily,
            "days": days,
            "needs": "Yes" if days < 30 else "No",
        })
    plan = sorted((row for row in stock if row["needs"] == "Yes"), key=lambda row: (row["days"], row["id"]))
    demand_expected = {str(row["Category"]): (int(row["Total_Units_Ordered"]), int(row["Order_Count"])) for row in demand_csv}
    demand_recomputed = {category: (demand_units[category], demand_lines[category]) for category in demand_units}
    return stock, plan, catalog, demand_expected, demand_recomputed, snapshot


def recalc(path: Path) -> Path | None:
    binary = shutil.which("libreoffice") or shutil.which("soffice")
    if not binary:
        return None
    root = Path(tempfile.mkdtemp(prefix="cowork-inventory-recalc-"))
    source_dir, output_dir = root / "source", root / "output"
    source_dir.mkdir(); output_dir.mkdir()
    copy = source_dir / OUTPUT
    shutil.copy2(path, copy)
    process = subprocess.run(
        [binary, "--headless", f"-env:UserInstallation={(root / 'profile').resolve().as_uri()}", "--convert-to", "xlsx", "--outdir", str(output_dir), str(copy)],
        stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300, check=False,
    )
    result = output_dir / OUTPUT
    return result if process.returncode == 0 and result.is_file() else None


def header(ws) -> list[str]:
    return [str(cell.value).strip() if cell.value is not None else "" for cell in ws[1]]


def rows(ws) -> list[tuple]:
    return [tuple(row) for row in ws.iter_rows(min_row=2, values_only=True) if any(value is not None and str(value).strip() for value in row)]


def evaluate(artifact_dir: Path, render: bool) -> dict:
    source = Path(__file__).resolve().parents[1] / "data" / "input_files"
    output = artifact_dir / OUTPUT
    results: list[dict] = []
    names = sorted(path.name for path in artifact_dir.iterdir()) if artifact_dir.is_dir() else []
    valid = output.is_file() and zipfile.is_zipfile(output)
    formula_wb = None
    cached_wb = None
    if valid:
        try:
            formula_wb = load_workbook(output, data_only=False)
            valid = formula_wb.sheetnames == SHEETS
        except Exception:
            valid = False
    results.append(criterion("R001", valid and names == [OUTPUT], f"valid={valid}; files={names}; sheets={formula_wb.sheetnames if formula_wb else None}"))
    if not valid:
        for cid in ("R002", "R003", "R004", "R005", "R006", "R007", "R008", "R009"):
            results.append(criterion(cid, False, "invalid workbook"))
        hashes_ok = all(digest(source / name) == value for name, value in INPUT_HASHES.items())
        results.append(criterion("R010", hashes_ok, f"input_sha256={hashes_ok}"))
        results.append(criterion("R011", True, "invalid or missing output triggers penalty"))
        return finalize(results, output)

    recalculated = recalc(output)
    if recalculated:
        cached_wb = load_workbook(recalculated, data_only=True)
    # When LibreOffice is unavailable, structural checks still run and numeric
    # checks use formula workbook values where possible.
    value_wb = cached_wb or formula_wb
    stock_expected, plan_expected, catalog_expected, demand_csv_expected, demand_recomputed, snapshot = expected_records(source)

    stock_ws, supplier_ws, demand_ws, plan_ws, summary_ws = (formula_wb[name] for name in SHEETS)
    stock_values, supplier_values, demand_values, plan_values, summary_values = (value_wb[name] for name in SHEETS)

    stock_headers = header(stock_ws)
    expected_stock_headers = ["Product_ID", "Product", "Category", "Current_Stock", "Total_Sales", "Daily_Rate", "Days_Remaining", "Needs_Restock"]
    stock_rows = rows(stock_values)
    stock_ok = stock_headers == expected_stock_headers and len(stock_rows) == len(stock_expected)
    if stock_ok:
        for actual, expected in zip(stock_rows, stock_expected):
            stock_ok &= actual[0] == expected["id"] and actual[1] == expected["name"] and normalize(actual[2]) == normalize(expected["category"])
            stock_ok &= number(actual[3]) == expected["stock"] and number(actual[4]) == expected["sales"]
    results.append(criterion("R002", stock_ok, f"Stock_Status rows={len(stock_rows)}/{len(stock_expected)}; source fields exact={stock_ok}"))

    formula_ok = stock_headers == expected_stock_headers
    if formula_ok:
        formula_ok = all(
            isinstance(stock_ws.cell(row, col).value, str) and stock_ws.cell(row, col).value.startswith("=")
            for row in range(2, len(stock_expected) + 2) for col in (6, 7, 8)
        )
    calculation_ok = formula_ok and len(stock_rows) == len(stock_expected)
    bad_calcs = []
    if calculation_ok:
        for actual, expected in zip(stock_rows, stock_expected):
            daily_value = number(actual[5])
            if daily_value is None or abs(round(daily_value, 2) - expected["daily"]) > 0.001:
                bad_calcs.append((expected["id"], "daily", actual[5], expected["daily"]))
            if number(actual[6]) is None or abs(number(actual[6]) - expected["days"]) > 0.11:
                bad_calcs.append((expected["id"], "days", actual[6], expected["days"]))
            if str(actual[7]).strip() != expected["needs"]:
                bad_calcs.append((expected["id"], "needs", actual[7], expected["needs"]))
    calculation_ok &= not bad_calcs
    results.append(criterion("R003", calculation_ok, f"native formulas={formula_ok}; calculation errors={bad_calcs[:5]}"))

    supplier_headers = header(supplier_ws)
    supplier_rows = rows(supplier_values)
    supplier_ok = supplier_headers == ["Supplier", "City", "Lead_Time_Days", "Min_Order_Qty", "Reliability_Score", "Risk_Index", "Risk_Band"] and len(supplier_rows) == len(catalog_expected)
    supplier_formula_ok = supplier_headers == ["Supplier", "City", "Lead_Time_Days", "Min_Order_Qty", "Reliability_Score", "Risk_Index", "Risk_Band"] and all(isinstance(supplier_ws.cell(row, 6).value, str) and supplier_ws.cell(row, 6).value.startswith("=") and isinstance(supplier_ws.cell(row, 7).value, str) and supplier_ws.cell(row, 7).value.startswith("=") for row in range(2, len(catalog_expected) + 2))
    if supplier_ok:
        for actual, expected in zip(supplier_rows, catalog_expected):
            lead, reliability = int(expected["lead_time_days"]), float(expected["reliability_score"])
            band = "High" if lead >= 14 or reliability < 4.0 else "Medium" if lead >= 10 or reliability < 4.5 else "Low"
            supplier_ok &= actual[0] == expected["name"] and actual[1] == expected.get("city", "")
            supplier_ok &= number(actual[2]) == lead and number(actual[3]) == int(expected["min_order_qty"]) and number(actual[4]) == reliability
            supplier_ok &= number(actual[5]) is not None and abs(number(actual[5]) - round(lead * (5 - reliability), 2)) <= 0.01 and str(actual[6]) == band
    results.append(criterion("R004", supplier_ok and supplier_formula_ok, f"supplier facts/risk={supplier_ok}; native risk formulas={supplier_formula_ok}"))

    demand_headers = header(demand_ws)
    demand_rows = rows(demand_values)
    demand_ok = demand_headers == ["Category", "Total_Units_Ordered", "Order_Count"] and len(demand_rows) == len(demand_csv_expected)
    if demand_ok:
        actual_map = {str(row[0]): (number(row[1]), number(row[2])) for row in demand_rows}
        demand_ok = all(category in actual_map and actual_map[category] == (float(values[0]), float(values[1])) for category, values in demand_csv_expected.items())
        demand_ok &= all(category in actual_map and actual_map[category] == (float(values[0]), float(values[1])) for category, values in demand_recomputed.items())
    results.append(criterion("R005", demand_ok, f"CSV and independent order/category aggregation match={demand_ok}; rows={len(demand_rows)}"))

    plan_headers = header(plan_ws)
    plan_rows = rows(plan_values)
    plan_formula_ok = plan_headers == ["Priority_Rank", "Product_ID", "Product", "Category", "Current_Stock", "Total_Sales", "Days_Remaining", "Recommended_Action"] and len(plan_rows) == len(plan_expected) and all(isinstance(plan_ws.cell(row, col).value, str) and plan_ws.cell(row, col).value.startswith("=") for row in range(2, len(plan_expected) + 2) for col in range(1, 8))
    plan_ok = plan_formula_ok
    if plan_ok:
        for rank, (actual, expected) in enumerate(zip(plan_rows, plan_expected), 1):
            plan_ok &= number(actual[0]) == rank and number(actual[1]) == expected["id"] and actual[2] == expected["name"] and normalize(actual[3]) == normalize(expected["category"])
            plan_ok &= number(actual[4]) == expected["stock"] and number(actual[5]) == expected["sales"] and number(actual[6]) is not None and abs(number(actual[6]) - expected["days"]) <= 0.11 and bool(str(actual[7]).strip())
    results.append(criterion("R006", plan_ok, f"restock rows={len(plan_rows)}/{len(plan_expected)}; linked facts and ordering={plan_ok}"))

    summary_headers = header(summary_ws)
    summary_rows = rows(summary_values)
    summary_map = {str(row[0]): row[1] for row in summary_rows if row and row[0] is not None}
    summary_formula_ok = summary_headers == ["Metric", "Value"] and all(isinstance(summary_ws.cell(row, 2).value, str) and summary_ws.cell(row, 2).value.startswith("=") for row in range(2, 6))
    summary_ok = summary_formula_ok and summary_map.get("Snapshot_Date") == snapshot
    expected_summary = {"Total_Products_Analyzed": len(stock_expected), "Products_Need_Restock": len(plan_expected), "Products_Healthy": len(stock_expected) - len(plan_expected), "Avg_Days_Remaining": round(sum(row["days"] for row in stock_expected) / len(stock_expected), 1)}
    for key, expected in expected_summary.items():
        actual = number(summary_map.get(key))
        summary_ok &= actual is not None and abs(actual - expected) <= 0.11
    results.append(criterion("R007", summary_ok, f"formula summary and independent values={summary_ok}"))

    structure_ok = formula_wb.sheetnames == SHEETS and all(formula_wb[name].freeze_panes == "A2" for name in SHEETS) and all(bool(formula_wb[name].tables) for name in SHEETS) and getattr(formula_wb.calculation, "calcMode", None) == "auto" and bool(getattr(formula_wb.calculation, "fullCalcOnLoad", False)) and bool(getattr(formula_wb.calculation, "forceFullCalc", False))
    results.append(criterion("R008", structure_ok, f"sheets/tables/freeze/recalc settings={structure_ok}"))

    roundtrip_ok = recalculated is not None and cached_wb is not None and cached_wb.sheetnames == SHEETS
    if render and roundtrip_ok:
        binary = shutil.which("libreoffice") or shutil.which("soffice")
        with tempfile.TemporaryDirectory(prefix="cowork-inventory-render-") as render_dir:
            profile = Path(render_dir) / "profile"
            process = subprocess.run([binary, "--headless", f"-env:UserInstallation={profile.resolve().as_uri()}", "--convert-to", "pdf", "--outdir", render_dir, str(output)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=300, check=False)
            roundtrip_ok &= process.returncode == 0 and (Path(render_dir) / (output.stem + ".pdf")).is_file()
    results.append(criterion("R009", roundtrip_ok, f"LibreOffice recalculation/render={roundtrip_ok}"))

    hashes_ok = all(digest(source / name) == value for name, value in INPUT_HASHES.items())
    results.append(criterion("R010", hashes_ok, f"input_sha256={hashes_ok}"))
    hardcoded_or_unchanged = digest(output) in {digest(source / name) for name in INPUT_HASHES} or not formula_ok or not summary_formula_ok
    results.append(criterion("R011", hardcoded_or_unchanged, "hardcoded/unchanged-input penalty triggered" if hardcoded_or_unchanged else "no hardcoded or unchanged-input trigger"))
    if cached_wb:
        cached_wb.close()
    if formula_wb:
        formula_wb.close()
    return finalize(results, output)


def finalize(results: list[dict], output: Path) -> dict:
    positive = sum(weight for weight in WEIGHTS.values() if weight > 0)
    earned = sum(item["score"] * WEIGHTS[item["criterion_id"]] for item in results if WEIGHTS[item["criterion_id"]] > 0)
    penalties = sum(WEIGHTS[item["criterion_id"]] for item in results if WEIGHTS[item["criterion_id"]] < 0 and item["score"] >= 1)
    reward = max(0.0, min(1.0, (earned + penalties) / positive))
    failures = [item["criterion_id"] for item in results if KINDS[item["criterion_id"]] == "hurdle" and item["score"] < 1]
    if failures:
        reward = min(reward, min(CAPS[cid] for cid in failures))
    return {
        "task_id": "cowork-evo/cowork-inventory-optimization",
        "rubric_version": "cowork-inventory-optimization-v0.1",
        "judge_version": "2.0.0",
        "artifact_sha256": [digest(output)] if output.is_file() else [],
        "criterion_results": [{"criterion_id": item["criterion_id"], "score": item["score"], "evidence": item["evidence"], "method": item["method"]} for item in sorted(results, key=lambda item: item["criterion_id"])],
        "reward": round(reward, 6),
        "passed": reward >= 0.8 and not failures,
        "hurdle_failures": failures,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    parser.add_argument("--skip-render", action="store_true")
    args = parser.parse_args()
    result = evaluate(args.artifact_dir.resolve(), not args.skip_render)
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.reward_file.write_text(f"{result['reward']:.6f}\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if result["passed"] else 1)


if __name__ == "__main__":
    main()
