#!/usr/bin/env python3
"""FORTE Finance-018 三期往来款工作簿整合器。"""

from __future__ import annotations

import argparse
import csv
import hashlib
import shutil
import sys
from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Iterable

sys.dont_write_bytecode = True

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
try:
    from pypinyin import lazy_pinyin
except ModuleNotFoundError:
    def lazy_pinyin(value: str) -> list[str]:
        """作者侧无 pypinyin 时的稳定降级；Harbor 镜像会使用正式依赖。"""
        return list(value)


PERIOD_FILES = (
    ("2025上半年", "2025往来明细-上半年.xlsx"),
    ("2025下半年", "2025往来明细-下半年.xlsx"),
    ("2026", "2026往来明细.xlsx"),
)
EXPECTED_HEADERS = (
    "科目名称",
    "客商名称",
    "方向",
    "期初余额",
    "本期借方",
    "本期贷方",
    "借方累计",
    "贷方累计",
    "方向",
    "期末余额",
)
CENT = Decimal("0.01")


@dataclass(frozen=True)
class Record:
    period: str
    source_file: str
    source_row: int
    subject: str
    counterparty: str
    opening_direction: str
    opening_balance: Decimal | None
    debit: Decimal | None
    credit: Decimal | None
    debit_total: Decimal | None
    credit_total: Decimal | None
    closing_direction: str
    closing_balance: Decimal | None
    closing_number_format: str


def _decimal(value: object) -> Decimal | None:
    if value in (None, ""):
        return None
    return Decimal(str(value)).quantize(CENT, rounding=ROUND_HALF_UP)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_records(input_dir: Path) -> tuple[list[Record], list[dict[str, object]]]:
    records: list[Record] = []
    sources: list[dict[str, object]] = []
    for period, filename in PERIOD_FILES:
        path = input_dir / filename
        if not path.is_file():
            raise FileNotFoundError(f"缺少上游指定输入：{filename}")
        workbook = load_workbook(path, data_only=True, read_only=False)
        if len(workbook.sheetnames) != 1:
            raise ValueError(f"{filename} 应只含一张明细表")
        sheet = workbook.active
        headers = tuple(cell.value for cell in sheet[1][:10])
        if headers != EXPECTED_HEADERS:
            raise ValueError(f"{filename} 表头与 Finance-018 原始格式不一致")
        count = 0
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, max_col=10), 2):
            values = [cell.value for cell in row]
            if not any(value is not None for value in values):
                continue
            if not values[0] or not values[1]:
                raise ValueError(f"{filename} 第 {row_number} 行缺少科目或客商")
            records.append(
                Record(
                    period=period,
                    source_file=filename,
                    source_row=row_number,
                    subject=str(values[0]),
                    counterparty=str(values[1]),
                    opening_direction=str(values[2] or ""),
                    opening_balance=_decimal(values[3]),
                    debit=_decimal(values[4]),
                    credit=_decimal(values[5]),
                    debit_total=_decimal(values[6]),
                    credit_total=_decimal(values[7]),
                    closing_direction=str(values[8] or ""),
                    closing_balance=_decimal(values[9]),
                    closing_number_format=row[9].number_format,
                )
            )
            count += 1
        sources.append(
            {
                "period": period,
                "filename": filename,
                "records": count,
                "size_bytes": path.stat().st_size,
                "sha256": _sha256(path),
            }
        )
    return records, sources


def _pinyin_key(record: Record) -> tuple[tuple[str, ...], Decimal, str]:
    amount = record.closing_balance or Decimal("0")
    return tuple(lazy_pinyin(record.counterparty)), -amount, record.subject


def current_balances(records: Iterable[Record]) -> tuple[list[Record], list[Record]]:
    current = [record for record in records if record.period == "2026"]
    unpaid = sorted(
        [record for record in current if record.closing_direction == "贷" and record.closing_balance not in (None, 0)],
        key=_pinyin_key,
    )
    unreceived = sorted(
        [record for record in current if record.closing_direction == "借" and record.closing_balance not in (None, 0)],
        key=_pinyin_key,
    )
    return unpaid, unreceived


def zombie_audit(records: Iterable[Record], unreceived: Iterable[Record]) -> list[dict[str, object]]:
    by_key: dict[tuple[str, str], dict[str, Record]] = {}
    for record in records:
        by_key.setdefault((record.subject, record.counterparty), {})[record.period] = record
    result: list[dict[str, object]] = []
    periods = [period for period, _ in PERIOD_FILES]
    for current in unreceived:
        history = by_key[(current.subject, current.counterparty)]
        balances = [history.get(period).closing_balance if history.get(period) else None for period in periods]
        directions = [history.get(period).closing_direction if history.get(period) else None for period in periods]
        is_zombie = (
            all(balance not in (None, Decimal("0")) for balance in balances)
            and all(direction == "借" for direction in directions)
            and len(set(balances)) == 1
        )
        result.append(
            {
                "subject": current.subject,
                "counterparty": current.counterparty,
                "balances": balances,
                "directions": directions,
                "is_zombie": is_zombie,
            }
        )
    return result


def _format_upstream_amount(record: Record) -> str:
    amount = record.closing_balance
    if amount is None:
        return ""
    if record.closing_number_format == "General" and amount == amount.to_integral_value():
        return str(int(amount))
    return f"{amount:,.2f}"


def write_upstream_csv(path: Path, records: Iterable[Record], amount_header: str) -> None:
    with path.open("w", encoding="gb18030", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\r\n")
        writer.writerow(["科目名称", "客商名称", amount_header])
        for record in records:
            writer.writerow([record.subject, record.counterparty, _format_upstream_amount(record)])


def _display_amount(value: Decimal | None) -> float | None:
    return None if value is None else float(value)


def _style_sheet(sheet, widths: list[float], landscape: bool = True) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D9E2F3")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = Border(bottom=thin_gray)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:1"
    sheet.print_area = sheet.dimensions
    sheet.oddFooter.center.text = "第 &P / &N 页"


def _add_table(sheet, name: str) -> None:
    if sheet.max_row < 2:
        return
    table = Table(displayName=name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def write_workbook(
    path: Path,
    records: list[Record],
    unpaid: list[Record],
    unreceived: list[Record],
    audit: list[dict[str, object]],
    sources: list[dict[str, object]],
) -> None:
    workbook = Workbook()
    detail = workbook.active
    detail.title = "跨期明细"
    detail.append(
        [
            "期间",
            "来源文件",
            "来源行",
            "科目名称",
            "客商名称",
            "期初方向",
            "期初余额",
            "本期借方",
            "本期贷方",
            "借方累计",
            "贷方累计",
            "期末方向",
            "期末余额",
        ]
    )
    for record in records:
        detail.append(
            [
                record.period,
                record.source_file,
                record.source_row,
                record.subject,
                record.counterparty,
                record.opening_direction,
                _display_amount(record.opening_balance),
                _display_amount(record.debit),
                _display_amount(record.credit),
                _display_amount(record.debit_total),
                _display_amount(record.credit_total),
                record.closing_direction,
                _display_amount(record.closing_balance),
            ]
        )
    for row in detail.iter_rows(min_row=2, min_col=7, max_col=13):
        for cell in row:
            cell.number_format = '#,##0.00;[Red]-#,##0.00'
    _style_sheet(detail, [14, 25, 10, 28, 34, 11, 15, 15, 15, 15, 15, 11, 15])
    _add_table(detail, "CrossPeriodDetail")

    for title, data, amount_header, table_name in (
        ("未付统计", unpaid, "未付款项", "UnpaidSummary"),
        ("未收统计", unreceived, "未收款项", "UnreceivedSummary"),
    ):
        sheet = workbook.create_sheet(title)
        sheet.append(["科目名称", "客商名称", amount_header])
        for record in data:
            sheet.append([record.subject, record.counterparty, _display_amount(record.closing_balance)])
        for cell in sheet["C"][1:]:
            cell.number_format = '#,##0.00;[Red]-#,##0.00'
        _style_sheet(sheet, [32, 38, 18], landscape=False)
        _add_table(sheet, table_name)

    zombie_sheet = workbook.create_sheet("僵尸账款核查")
    zombie_sheet.append(
        ["科目名称", "客商名称", "2025上半年", "2025下半年", "2026", "是否僵尸账"]
    )
    for item in audit:
        zombie_sheet.append(
            [
                item["subject"],
                item["counterparty"],
                *[_display_amount(value) for value in item["balances"]],
                "是" if item["is_zombie"] else "否",
            ]
        )
    for row in zombie_sheet.iter_rows(min_row=2, min_col=3, max_col=5):
        for cell in row:
            cell.number_format = '#,##0.00;[Red]-#,##0.00'
    _style_sheet(zombie_sheet, [32, 38, 18, 18, 18, 15])
    _add_table(zombie_sheet, "ZombieAudit")

    summary = workbook.create_sheet("校验摘要")
    summary.append(["校验项", "结果"])
    summary_rows = [
        ("原始记录总数", len(records)),
        ("2026 未付笔数", len(unpaid)),
        ("2026 未付总额", float(sum((r.closing_balance or Decimal("0")) for r in unpaid))),
        ("2026 未收笔数", len(unreceived)),
        ("2026 未收总额", float(sum((r.closing_balance or Decimal("0")) for r in unreceived))),
        ("僵尸账笔数", sum(bool(item["is_zombie"]) for item in audit)),
        ("僵尸账结论", "存在僵尸账款" if any(item["is_zombie"] for item in audit) else "无僵尸账款"),
    ]
    for item in sources:
        summary_rows.extend(
            [
                (f"{item['period']} 记录数", item["records"]),
                (f"{item['period']} 输入 SHA256", item["sha256"]),
            ]
        )
    for row in summary_rows:
        summary.append(list(row))
    for cell in summary["B"]:
        if isinstance(cell.value, float):
            cell.number_format = '#,##0.00;[Red]-#,##0.00'
    _style_sheet(summary, [30, 72], landscape=False)
    _add_table(summary, "ValidationSummary")

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(path)


def write_audit_text(path: Path, audit: list[dict[str, object]]) -> None:
    zombies = [item for item in audit if item["is_zombie"]]
    lines = [
        f"僵尸账款核查结论：{'存在僵尸账款' if zombies else '无僵尸账款'}。",
        "核查口径：仅核查 2026 年期末仍为借方余额的款项；同一客商、同一科目在三个期间均为非零借方余额且金额完全相同，才判定为僵尸账款。",
        "",
    ]
    if zombies:
        for item in zombies:
            amount = item["balances"][-1]
            lines.append(f"- {item['counterparty']} | {item['subject']} | {amount:,.2f}")
    else:
        lines.append("未发现满足上述三期连续不动条件的往来款。")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def generate(input_dir: Path, output_dir: Path, processor_source: Path | None = None) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    records, sources = read_records(input_dir)
    unpaid, unreceived = current_balances(records)
    audit = zombie_audit(records, unreceived)
    write_upstream_csv(output_dir / "未付统计.csv", unpaid, "未付款项")
    write_upstream_csv(output_dir / "未收统计.csv", unreceived, "未收款项")
    write_audit_text(output_dir / "僵尸账款核查.txt", audit)
    write_workbook(output_dir / "往来款跨期整合.xlsx", records, unpaid, unreceived, audit, sources)
    if processor_source is not None:
        destination = output_dir / "reconcile_workbooks.py"
        if processor_source.resolve() != destination.resolve():
            shutil.copy2(processor_source, destination)


def main() -> None:
    parser = argparse.ArgumentParser(description="整合 FORTE Finance-018 的三期往来明细")
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    generate(args.input_dir, args.output_dir, Path(__file__))


if __name__ == "__main__":
    main()
