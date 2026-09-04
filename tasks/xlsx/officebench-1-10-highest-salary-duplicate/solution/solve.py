from pathlib import Path
import os
from openpyxl import load_workbook
INPUT = 'salary.xlsx'
OUTPUT = 'salary_highest_duplicate.xlsx'
OPERATION = 'highest_salary_duplicate'
def main():
    root = Path(__file__).parents[1]
    src = root / "data/input_files" / INPUT
    out_dir = Path(os.environ.get("OUTPUT_DIR", "/app/output"))
    out_dir.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(src)
    ws = wb.active
    if OPERATION == "attendance":
        ws.cell(1, 2).value = "Attendance"
        for row in range(2, ws.max_row + 1):
            ws.cell(row, 2).value = "Y"
    elif OPERATION == "highest_salary_duplicate":
        amount, name = max((ws.cell(r, 2).value, ws.cell(r, 1).value) for r in range(2, ws.max_row + 1))
        ws.append([name, amount])
    elif OPERATION == "lowest_budget_add_highest":
        amounts = [ws.cell(r, 2).value for r in range(2, ws.max_row + 1)]
        ws.cell(amounts.index(max(amounts)) + 2, 2).value = max(amounts) + min(amounts)
    elif OPERATION in {"total_salary", "total_budget", "total_shopping_weight"}:
        ws.append(["Total", sum(ws.cell(r, 2).value for r in range(2, ws.max_row + 1))])
    else:
        raise ValueError(OPERATION)
    wb.save(out_dir / OUTPUT)
    wb.close()
if __name__ == "__main__":
    main()
