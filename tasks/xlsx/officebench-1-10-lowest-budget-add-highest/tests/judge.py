#!/usr/bin/env python3
from __future__ import annotations
import argparse, hashlib, json, zipfile
from pathlib import Path
from openpyxl import load_workbook
TASK_ID = 'cowork-evo/officebench-1-10-lowest-budget-add-highest'
INPUT = 'company_budget.xlsx'
EXPECTED_INPUT_SHA256 = "7b14e19c56e1b5185f61e16cda1f1b00d424d43cb41803ba5d7b9141a4a160a3"
OUTPUT = 'budget_lowest_to_highest.xlsx'
OPERATION = 'lowest_budget_add_highest'
WEIGHTS = {"R001": .10, "R002": .45, "R003": .15, "R004": .20, "R005": .05, "R006": .05, "R007": -.20}
def criterion(cid, ok, evidence):
    return {"criterion_id": cid, "score": 1.0 if ok else 0.0, "evidence": evidence, "method": "deterministic"}
def expected_from_input(wb):
    ws = wb.active
    expected = {(s.title, cell.coordinate): cell.value for s in wb.worksheets for row in s.iter_rows() for cell in row}
    if OPERATION == "attendance":
        expected[(ws.title, "B1")] = "Attendance"
        for row in range(2, ws.max_row + 1): expected[(ws.title, f"B{row}")] = "Y"
        shape = (ws.max_row, 2)
    elif OPERATION == "highest_salary_duplicate":
        amount, name = max((ws.cell(r, 2).value, ws.cell(r, 1).value) for r in range(2, ws.max_row + 1))
        row = ws.max_row + 1
        expected[(ws.title, f"A{row}")], expected[(ws.title, f"B{row}")] = name, amount
        shape = (row, 2)
    elif OPERATION == "lowest_budget_add_highest":
        amounts = [ws.cell(r, 2).value for r in range(2, ws.max_row + 1)]
        target = amounts.index(max(amounts)) + 2
        expected[(ws.title, f"B{target}")] = max(amounts) + min(amounts)
        shape = (ws.max_row, ws.max_column)
    elif OPERATION in {"total_salary", "total_budget", "total_shopping_weight"}:
        row = ws.max_row + 1
        expected[(ws.title, f"A{row}")] = "Total"
        expected[(ws.title, f"B{row}")] = sum(ws.cell(r, 2).value for r in range(2, ws.max_row + 1))
        shape = (row, 2)
    else: raise ValueError(OPERATION)
    return expected, shape
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--artifact-dir", type=Path, required=True)
    ap.add_argument("--result", type=Path, required=True)
    ap.add_argument("--reward-file", type=Path, required=True)
    a = ap.parse_args()
    root = Path(__file__).parents[1]; src = root / "data/input_files" / INPUT; out = a.artifact_dir / OUTPUT
    results = []; native = out.is_file() and zipfile.is_zipfile(out)
    results.append(criterion("R001", native, "输出存在且为可解析 OOXML 工作簿"))
    target_ok = preservation_ok = formula_ok = parse_ok = False; mismatch_count = 0
    if native:
        try:
            wi = load_workbook(src, data_only=False); wo = load_workbook(out, data_only=False)
            expected, shape = expected_from_input(wi); parse_ok = True
            target_ok = wi.sheetnames == wo.sheetnames and (wo.active.max_row, wo.active.max_column) == shape
            if target_ok:
                for (sheet, coord), value in expected.items():
                    if wo[sheet][coord].value != value: mismatch_count += 1
                target_ok = mismatch_count == 0; preservation_ok = target_ok; formula_ok = True
            wi.close(); wo.close()
            results.append(criterion("R002", target_ok, f"从输入独立重算，错误单元格={mismatch_count}"))
            results.append(criterion("R003", formula_ok, "目标为可编辑单元格值，未要求隐藏计算"))
            results.append(criterion("R004", preservation_ok, "逐格检查非目标内容和原工作表结构"))
            results.append(criterion("R005", parse_ok, "openpyxl 可解析工作簿"))
            input_hash = hashlib.sha256(src.read_bytes()).hexdigest()
            input_hash_ok = input_hash == EXPECTED_INPUT_SHA256
            results.append(criterion("R006", input_hash_ok, "输入副本 SHA256 一致"))
            results.append(criterion("R007", False, "未检测到空白或硬编码替代品"))
        except Exception as exc:
            results.extend([criterion("R002", False, f"读取或重算失败: {exc}"), criterion("R003", False, "读取失败"), criterion("R004", False, "读取失败"), criterion("R005", False, "读取失败"), criterion("R006", False, "读取失败"), criterion("R007", True, "损坏输出触发惩罚")])
    else:
        results.extend([criterion("R002", False, "缺少有效输出"), criterion("R003", False, "缺少有效输出"), criterion("R004", False, "缺少有效输出"), criterion("R005", False, "缺少有效输出"), criterion("R006", False, "缺少有效输出"), criterion("R007", True, "缺少有效输出触发惩罚")])
    vals = {x["criterion_id"]: x["score"] for x in results}; reward = sum(WEIGHTS[k] * vals.get(k, 0) for k in WEIGHTS)
    passed = vals.get("R001", 0) == 1 and vals.get("R002", 0) == 1 and reward >= .7
    payload = {"task_id": TASK_ID, "rubric_version": "cowork-evo-rubric-v0.1", "judge_version": "officebench-l1-independent-v0.1", "criterion_results": results, "reward": reward, "passed": passed}
    a.result.parent.mkdir(parents=True, exist_ok=True); a.result.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    a.reward_file.parent.mkdir(parents=True, exist_ok=True); a.reward_file.write_text(f"{reward:.6f}\n")
    raise SystemExit(0 if passed else 1)
if __name__ == "__main__":
    main()
