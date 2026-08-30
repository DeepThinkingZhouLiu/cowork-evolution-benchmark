#!/usr/bin/env python3
"""确定性 Judge：按输入计算 midterm1 降序结果并检查交付。"""

from __future__ import annotations

import argparse
import hashlib
import json
import zipfile
from pathlib import Path

from openpyxl import load_workbook

TASK_ID = "cowork-evo/officebench-1-11-score-sort"
RUBRIC_VERSION = "officebench-1-11-score-sort-v0.1"
JUDGE_VERSION = "officebench-1-11-score-sort-deterministic-v0.1"
INPUT_NAME = "score_input.xlsx"
OUTPUT_NAME = "sorted_score.xlsx"
EXPECTED_INPUT_SHA256 = "b0b81274ef8087f31e6920d8ff7ae3358c082979533f71b2fde134075744c973"
WEIGHTS = {"R001": 0.20, "R002": 0.40, "R003": 0.20, "R004": 0.10, "R005": 0.05, "R006": 0.05}


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def criterion(cid: str, score: bool | float, evidence: str) -> dict:
    return {"criterion_id": cid, "score": float(score), "evidence": evidence, "method": "deterministic"}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    input_path = Path(__file__).parents[1] / "data/input_files" / INPUT_NAME
    output_path = args.artifact_dir / OUTPUT_NAME
    results: list[dict] = []
    try:
        if not output_path.is_file() or not zipfile.is_zipfile(output_path):
            raise ValueError("sorted_score.xlsx 缺失或不是有效 OOXML")
        source = load_workbook(input_path, data_only=False)
        candidate = load_workbook(output_path, data_only=False)
        sws, cws = source["Sheet1"], candidate["Sheet1"]
        source_rows = [[cell.value for cell in row] for row in sws.iter_rows(min_row=2)]
        expected_rows = sorted(source_rows, key=lambda row: row[1], reverse=True)
        actual_rows = [[cell.value for cell in row] for row in cws.iter_rows(min_row=2, max_row=cws.max_row, max_col=sws.max_column)]
        structure_ok = candidate.sheetnames == source.sheetnames and cws.max_row == sws.max_row and cws.max_column == sws.max_column
        results.append(criterion("R001", structure_ok, f"工作表={candidate.sheetnames}; 尺寸={cws.max_row}x{cws.max_column}"))
        header_ok = [cws.cell(1, col).value for col in range(1, sws.max_column + 1)] == [sws.cell(1, col).value for col in range(1, sws.max_column + 1)]
        order_ok = actual_rows == expected_rows
        results.append(criterion("R002", header_ok and order_ok, f"表头保持={header_ok}; 独立降序结果一致={order_ok}"))
        results.append(criterion("R003", sorted(actual_rows) == sorted(source_rows), f"数据行多重集合保持={sorted(actual_rows) == sorted(source_rows)}"))
        style_ok = all(cws.cell(row, col)._style == sws.cell(row, col)._style for row in range(1, sws.max_row + 1) for col in range(1, sws.max_column + 1))
        results.append(criterion("R004", style_ok, f"单元格样式保持={style_ok}"))
        results.append(criterion("R005", digest(input_path) == EXPECTED_INPUT_SHA256, f"输入 SHA256={digest(input_path)}"))
        results.append(criterion("R006", cws.protection.sheet is not True, "工作表可编辑"))
        source.close(); candidate.close()
    except Exception as exc:
        results = [criterion("R001", False, f"Judge 读取失败：{exc}")]
    by_id = {item["criterion_id"]: item["score"] for item in results}
    reward = sum(WEIGHTS[cid] * by_id.get(cid, 0.0) for cid in WEIGHTS)
    passed = reward >= 0.7 and all(by_id.get(cid) == 1.0 for cid in ("R001", "R002"))
    payload = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [digest(output_path)] if output_path.is_file() else [], "criterion_results": results, "reward": reward, "passed": passed}
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.write_text(f"{reward:.6f}\n", encoding="utf-8")
    raise SystemExit(0 if passed else 1)


if __name__ == "__main__":
    main()
