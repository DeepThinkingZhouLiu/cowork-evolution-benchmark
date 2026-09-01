#!/usr/bin/env python3
"""将成绩输入投影为一个含两个考试工作表的 Excel 工作簿。"""

from __future__ import annotations

import argparse
import hashlib
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook

INPUT_NAME = "score_input.xlsx"
OUTPUT_NAME = "midterm_split.xlsx"
EXPECTED_INPUT_SHA256 = "b0b81274ef8087f31e6920d8ff7ae3358c082979533f71b2fde134075744c973"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def copy_cell(source_cell, target_cell) -> None:
    target_cell.value = source_cell.value
    if source_cell.has_style:
        target_cell.font = copy(source_cell.font)
        target_cell.fill = copy(source_cell.fill)
        target_cell.border = copy(source_cell.border)
        target_cell.alignment = copy(source_cell.alignment)
        target_cell.protection = copy(source_cell.protection)
        target_cell.number_format = source_cell.number_format


def write_exam_sheet(workbook: Workbook, source_sheet, sheet_name: str, score_column: int) -> None:
    sheet = workbook.create_sheet(sheet_name)
    for row in range(1, source_sheet.max_row + 1):
        copy_cell(source_sheet.cell(row, 1), sheet.cell(row, 1))
        copy_cell(source_sheet.cell(row, score_column), sheet.cell(row, 2))
    score_letter = "B" if score_column == 2 else "C"
    sheet.column_dimensions["A"].width = source_sheet.column_dimensions["A"].width
    sheet.column_dimensions["B"].width = source_sheet.column_dimensions[score_letter].width
    sheet.sheet_view.showGridLines = source_sheet.sheet_view.showGridLines
    sheet.freeze_panes = source_sheet.freeze_panes


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    source = args.input_dir / INPUT_NAME
    if sha256(source) != EXPECTED_INPUT_SHA256:
        raise ValueError("score_input.xlsx 输入哈希不匹配")
    source_book = load_workbook(source, data_only=False)
    try:
        source_sheet = source_book["Sheet1"]
        if source_sheet.max_column < 3 or source_sheet.max_row < 1:
            raise ValueError("输入缺少 Name、midterm1 或 midterm2 列")
        workbook = Workbook()
        workbook.remove(workbook.active)
        write_exam_sheet(workbook, source_sheet, "midterm1", 2)
        write_exam_sheet(workbook, source_sheet, "midterm2", 3)
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        args.output_dir.mkdir(parents=True, exist_ok=True)
        workbook.save(args.output_dir / OUTPUT_NAME)
        workbook.close()
    finally:
        source_book.close()


if __name__ == "__main__":
    main()
