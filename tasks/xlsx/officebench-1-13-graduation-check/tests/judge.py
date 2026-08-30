#!/usr/bin/env python3
"""确定性 Judge：从课程工作簿重算 CS/MATH 数量及毕业结论。"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

TASK_ID = "cowork-evo/officebench-1-13-graduation-check"
RUBRIC_VERSION = "officebench-1-13-graduation-check-v0.1"
JUDGE_VERSION = "officebench-1-13-graduation-check-deterministic-v0.1"
INPUT_NAME = "Thomas_Robinson_Course_List.xlsx"
OUTPUT_NAME = "answer.txt"
EXPECTED_INPUT_SHA256 = "5a495d6f9d3c319207dc1b9ed002672a42eee30c3f25741b31ca4e7df399faeb"
WEIGHTS = {"R001": 0.20, "R002": 0.20, "R003": 0.20, "R004": 0.20, "R005": 0.15, "R006": 0.05}


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def c(cid: str, score: bool | float, evidence: str) -> dict:
    return {"criterion_id": cid, "score": float(score), "evidence": evidence, "method": "deterministic"}


def main() -> None:
    parser = argparse.ArgumentParser(); parser.add_argument("--artifact-dir", type=Path, required=True); parser.add_argument("--result", type=Path, required=True); parser.add_argument("--reward-file", type=Path, required=True); args = parser.parse_args()
    input_path = Path(__file__).parents[1] / "data/input_files" / INPUT_NAME
    output_path = args.artifact_dir / OUTPUT_NAME
    results: list[dict] = []
    try:
        if not output_path.is_file():
            raise ValueError("answer.txt 缺失")
        text = output_path.read_text(encoding="utf-8")
        answer = text.strip()
        source = load_workbook(input_path, data_only=True); sheet = source["Sheet1"]
        cs_count = sum(str(sheet.cell(row, 3).value or "").startswith("CS ") for row in range(2, sheet.max_row + 1))
        math_count = sum(str(sheet.cell(row, 3).value or "").startswith("MATH ") for row in range(2, sheet.max_row + 1))
        expected = "Yes" if cs_count >= 8 and math_count >= 5 else "No"
        results.append(c("R001", True, "answer.txt 为 UTF-8 可读文本"))
        lines = text.splitlines()
        results.append(c("R002", len(lines) == 1 and answer in {"Yes", "No"} and lines[0] == answer, f"原始文本={text!r}"))
        results.append(c("R003", cs_count == 5, f"独立统计 CS={cs_count}"))
        results.append(c("R004", math_count == 5, f"独立统计 MATH={math_count}"))
        results.append(c("R005", answer == expected and expected == "No", f"独立规则答案={expected}; 输出={answer!r}"))
        results.append(c("R006", digest(input_path) == EXPECTED_INPUT_SHA256, f"输入 SHA256={digest(input_path)}"))
        source.close()
    except Exception as exc:
        results = [c("R001", False, f"Judge 读取失败：{exc}")]
    by_id = {item["criterion_id"]: item["score"] for item in results}; reward = sum(WEIGHTS[key] * by_id.get(key, 0.0) for key in WEIGHTS)
    passed = reward >= 0.7 and all(by_id.get(key) == 1.0 for key in ("R001", "R002", "R005"))
    payload = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [digest(output_path)] if output_path.is_file() else [], "criterion_results": results, "reward": reward, "passed": passed}
    args.result.parent.mkdir(parents=True, exist_ok=True); args.result.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.parent.mkdir(parents=True, exist_ok=True); args.reward_file.write_text(f"{reward:.6f}\n", encoding="utf-8")
    raise SystemExit(0 if passed else 1)


if __name__ == "__main__": main()
