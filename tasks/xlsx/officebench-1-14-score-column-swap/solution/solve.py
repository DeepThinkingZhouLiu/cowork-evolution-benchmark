#!/usr/bin/env python3
"""OfficeBench 1-14/1 的逐题参考解法。"""

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

from openpyxl import load_workbook

INPUT_NAME = "score_input.xlsx"
OUTPUT_NAME = "score.xlsx"
EXPECTED_INPUT_SHA256 = "6b98c5b9d854ab3d5fce62632a58a49c58559fff04435bb59d68fe1f0c6dd255"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    source = args.input_dir / INPUT_NAME
    if sha256(source) != EXPECTED_INPUT_SHA256:
        raise ValueError("score.xlsx 输入哈希不匹配")
    workbook = load_workbook(source)
    sheet = workbook["Sheet1"]
    if [sheet.cell(1, c).value for c in range(1, 4)] != ["Name", "midterm1", "midterm2"]:
        raise ValueError("输入表头不在预期状态")
    for row in range(1, sheet.max_row + 1):
        sheet.cell(row, 2).value, sheet.cell(row, 3).value = sheet.cell(row, 3).value, sheet.cell(row, 2).value
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    args.output_dir.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output_dir / OUTPUT_NAME)
    workbook.close()


if __name__ == "__main__":
    main()
