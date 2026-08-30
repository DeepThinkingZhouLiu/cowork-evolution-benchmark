#!/usr/bin/env python3
"""OfficeBench 1-4/0 的逐题参考解法。"""

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

from openpyxl import load_workbook

INPUT_NAME = "score_input.xlsx"
OUTPUT_NAME = "score.xlsx"
EXPECTED_INPUT_SHA256 = "a5c8026dc09f38411c6b8783dad74f7f66f08387eeb21b9d8eb7f28a90dd5b6e"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    source = args.input_dir / INPUT_NAME
    if sha256(source) != EXPECTED_INPUT_SHA256:
        raise ValueError("score.xlsx 输入哈希不匹配")
    workbook = load_workbook(source)
    if workbook.sheetnames != ["Sheet1"]:
        raise ValueError(f"未知工作表结构：{workbook.sheetnames}")
    sheet = workbook["Sheet1"]
    if sheet["A12"].value != "Bob":
        raise ValueError("输入中的 Bob 不在预期行")
    sheet["B12"] = 100
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    args.output_dir.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output_dir / OUTPUT_NAME)
    workbook.close()


if __name__ == "__main__":
    main()
