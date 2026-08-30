#!/usr/bin/env python3
"""确定性 Judge：检查单值修改、非目标保留和原生工作簿完整性。"""

from __future__ import annotations

import argparse
import hashlib
import json
import zipfile
from pathlib import Path

from openpyxl import load_workbook

TASK_ID = "cowork-evo/officebench-1-4-score-update"
RUBRIC_VERSION = "officebench-1-4-v0.1"
JUDGE_VERSION = "officebench-1-4-deterministic-v0.1"
OUTPUT_NAME = "score.xlsx"
INPUT_NAME = "score_input.xlsx"
EXPECTED_INPUT_SHA256 = "a5c8026dc09f38411c6b8783dad74f7f66f08387eeb21b9d8eb7f28a90dd5b6e"


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def criterion(cid: str, score: float, evidence: str) -> dict:
    return {"criterion_id": cid, "score": score, "evidence": evidence, "method": "deterministic"}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    input_path = Path(__file__).parents[1] / "data/input_files" / INPUT_NAME
    out = args.artifact_dir / OUTPUT_NAME
    results = []
    scores = []
    if not out.is_file():
        results = [criterion("R001", 0, "score.xlsx 缺失")]
        reward = 0.0
    else:
        try:
            if not zipfile.is_zipfile(out):
                raise ValueError("不是 OOXML ZIP")
            source = load_workbook(input_path, data_only=False)
            candidate = load_workbook(out, data_only=False)
            same_structure = candidate.sheetnames == source.sheetnames and candidate["Sheet1"].max_row == source["Sheet1"].max_row and candidate["Sheet1"].max_column == source["Sheet1"].max_column
            results.append(criterion("R001", 1.0 if same_structure else 0.0, f"工作表结构={candidate.sheetnames}, 尺寸={candidate['Sheet1'].max_row}x{candidate['Sheet1'].max_column}"))
            results.append(criterion("R002", 1.0 if candidate["Sheet1"]["A12"].value == "Bob" and candidate["Sheet1"]["B12"].value == 100 else 0.0, f"Sheet1!B12={candidate['Sheet1']['B12'].value!r}"))
            changed = []
            for row in range(1, source["Sheet1"].max_row + 1):
                for col in range(1, source["Sheet1"].max_column + 1):
                    coord = source["Sheet1"].cell(row, col).coordinate
                    if coord == "B12":
                        continue
                    if source["Sheet1"].cell(row, col).value != candidate["Sheet1"].cell(row, col).value:
                        changed.append(coord)
            results.append(criterion("R003", 1.0 if not changed else 0.0, f"非目标变更={changed[:12]}"))
            style_ok = all(source["Sheet1"].cell(row, col)._style == candidate["Sheet1"].cell(row, col)._style for row in range(1, source["Sheet1"].max_row + 1) for col in range(1, source["Sheet1"].max_column + 1) if source["Sheet1"].cell(row, col).coordinate != "B12")
            results.append(criterion("R004", 1.0 if style_ok else 0.0, f"非目标样式保持={style_ok}"))
            results.append(criterion("R005", 1.0 if digest(input_path) == EXPECTED_INPUT_SHA256 else 0.0, f"输入 SHA256={digest(input_path)}"))
            results.append(criterion("R006", 1.0 if candidate["Sheet1"].protection.sheet is not True else 0.0, "工作表可编辑"))
            source.close(); candidate.close()
            scores = [x["score"] for x in results]
            reward = sum(scores) / len(scores)
        except Exception as exc:
            results.append(criterion("R001", 0, f"读取失败：{exc}"))
            reward = 0.0
    passed = reward >= 0.7 and all(x["score"] == 1.0 for x in results[:2])
    payload = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [digest(out)] if out.is_file() else [], "criterion_results": results, "reward": reward, "passed": passed}
    args.result.parent.mkdir(parents=True, exist_ok=True); args.result.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.parent.mkdir(parents=True, exist_ok=True); args.reward_file.write_text(f"{reward:.6f}\n", encoding="utf-8")
    raise SystemExit(0 if passed else 1)


if __name__ == "__main__":
    main()
