#!/usr/bin/env python3
"""OfficeBench 1-5/3 的逐题参考解法。"""

from __future__ import annotations

import argparse
import hashlib
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

INPUT_NAME = "company_budget_input.xlsx"
OUTPUT_NAME = "company_budget.xlsx"
EXPECTED_SHA256 = "7b14e19c56e1b5185f61e16cda1f1b00d424d43cb41803ba5d7b9141a4a160a3"


def digest(path: Path) -> str:
    value = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            value.update(block)
    return value.hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    source = args.input_dir / INPUT_NAME
    if digest(source) != EXPECTED_SHA256:
        raise ValueError("输入工作簿哈希不匹配")
    workbook = load_workbook(source)
    if workbook.sheetnames != ["Sheet1"]:
        raise ValueError(f"未知工作表结构：{workbook.sheetnames}")
    sheet = workbook["Sheet1"]
    expected_rows = [("salary", 2000000), ("computer", 50000), ("entertain", 1000)]
    if list(sheet.iter_rows(min_row=2, max_row=4, values_only=True)) != expected_rows:
        raise ValueError("预算输入值与审计证据不一致")
    sheet["A5"] = "predicted budget"
    sheet["B5"] = "=SUM(B2:B4)*(1+10%)^3"
    for column in (1, 2):
        sheet.cell(5, column)._style = copy(sheet.cell(4, column)._style)
    sheet["B5"].number_format = "#,##0"
    sheet.column_dimensions["A"].width = max(sheet.column_dimensions["A"].width or 0, 20)
    sheet.column_dimensions["B"].width = max(sheet.column_dimensions["B"].width or 0, 16)
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    args.output_dir.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output_dir / OUTPUT_NAME)
    workbook.close()


if __name__ == "__main__":
    main()
