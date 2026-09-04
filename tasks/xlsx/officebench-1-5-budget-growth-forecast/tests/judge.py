#!/usr/bin/env python3
"""确定性 Judge：从输入重算复利预测并验证公式的变异行为。"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from openpyxl import load_workbook

TASK_ID = "cowork-evo/officebench-1-5-budget-growth-forecast"
RUBRIC_VERSION = "officebench-1-5-budget-growth-forecast-v0.1"
JUDGE_VERSION = "officebench-1-5-budget-growth-forecast-deterministic-v0.1"
INPUT_NAME = "company_budget_input.xlsx"
OUTPUT_NAME = "company_budget.xlsx"
EXPECTED_SHA256 = "7b14e19c56e1b5185f61e16cda1f1b00d424d43cb41803ba5d7b9141a4a160a3"
WEIGHTS = {"R001": 0.15, "R002": 0.15, "R003": 0.25, "R004": 0.15, "R005": 0.15, "R006": 0.05, "R007": 0.05, "R008": 0.05}


def digest(path: Path) -> str:
    value = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            value.update(block)
    return value.hexdigest()


def criterion(cid: str, score: bool | float, evidence: str) -> dict:
    return {"criterion_id": cid, "score": float(score), "evidence": evidence, "method": "deterministic"}


def recalculated_value(path: Path) -> float | None:
    with tempfile.TemporaryDirectory(prefix="xlsx-budget-recalc-") as raw:
        root = Path(raw)
        source_dir = root / "source"
        output_dir = root / "output"
        profile = root / "profile"
        source_dir.mkdir()
        output_dir.mkdir()
        profile.mkdir()
        candidate = source_dir / OUTPUT_NAME
        shutil.copy2(path, candidate)
        command = [
            "libreoffice",
            f"-env:UserInstallation={profile.resolve().as_uri()}",
            "--headless",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(output_dir),
            str(candidate),
        ]
        subprocess.run(command, check=True, capture_output=True, text=True, timeout=90)
        recalculated = output_dir / OUTPUT_NAME
        workbook = load_workbook(recalculated, data_only=True)
        value = workbook["Sheet1"]["B5"].value
        workbook.close()
        try:
            return float(value)
        except (TypeError, ValueError):
            return None


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    input_path = Path(__file__).parents[1] / "data/input_files" / INPUT_NAME
    output_path = args.artifact_dir / OUTPUT_NAME
    results: list[dict] = []
    try:
        if not output_path.is_file() or not zipfile.is_zipfile(output_path):
            raise ValueError("company_budget.xlsx 缺失或不是有效 OOXML")
        source = load_workbook(input_path, data_only=False)
        candidate = load_workbook(output_path, data_only=False)
        structure_ok = candidate.sheetnames == source.sheetnames == ["Sheet1"] and candidate["Sheet1"].max_row == 5 and candidate["Sheet1"].max_column == 2
        results.append(criterion("R001", structure_ok, f"工作表={candidate.sheetnames}; 尺寸={candidate['Sheet1'].max_row}x{candidate['Sheet1'].max_column}"))
        results.append(criterion("R002", candidate["Sheet1"]["A5"].value == "predicted budget", f"A5={candidate['Sheet1']['A5'].value!r}"))
        formula = candidate["Sheet1"]["B5"].value
        formula_ok = isinstance(formula, str) and formula.startswith("=")
        value = recalculated_value(output_path) if formula_ok else formula if isinstance(formula, (int, float)) else None
        expected = sum(source["Sheet1"].cell(row, 2).value for row in range(2, 5)) * 1.10**3
        results.append(criterion("R003", value is not None and abs(float(value) - expected) < 0.01, f"输出值={value}; 输入重算真值={expected}"))
        original_ok = all(source["Sheet1"].cell(row, col).value == candidate["Sheet1"].cell(row, col).value for row in range(1, 5) for col in range(1, 3))
        results.append(criterion("R004", original_ok, f"原始 4x2 区域保持={original_ok}"))
        mutation_ok = False
        mutation_value = None
        if formula_ok:
            with tempfile.TemporaryDirectory(prefix="xlsx-budget-mutation-") as raw:
                mutated = Path(raw) / OUTPUT_NAME
                candidate["Sheet1"]["B2"] = float(candidate["Sheet1"]["B2"].value) + 1000
                candidate.save(mutated)
                mutation_value = recalculated_value(mutated)
                mutation_expected = (expected / 1.10**3 + 1000) * 1.10**3
                mutation_ok = mutation_value is not None and abs(mutation_value - mutation_expected) < 0.01
        results.append(criterion("R005", formula_ok and mutation_ok, f"公式={formula!r}; 变异重算值={mutation_value}"))
        presentation_ok = candidate["Sheet1"]["B5"].number_format != "General" and (candidate["Sheet1"].column_dimensions["A"].width or 0) >= 16
        results.append(criterion("R006", presentation_ok, f"数字格式={candidate['Sheet1']['B5'].number_format}; A列宽={candidate['Sheet1'].column_dimensions['A'].width}"))
        results.append(criterion("R007", candidate["Sheet1"].protection.sheet is not True, "工作表可编辑"))
        results.append(criterion("R008", digest(input_path) == EXPECTED_SHA256, f"输入 SHA256={digest(input_path)}"))
        source.close()
        candidate.close()
    except Exception as exc:
        results = [criterion("R001", False, f"Judge 读取或重算失败：{exc}")]
    result_by_id = {item["criterion_id"]: item["score"] for item in results}
    reward = sum(WEIGHTS[cid] * result_by_id.get(cid, 0.0) for cid in WEIGHTS)
    hurdle_ids = ("R001", "R002", "R003", "R004", "R005")
    passed = reward >= 0.8 and all(result_by_id.get(cid) == 1.0 for cid in hurdle_ids)
    payload = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [digest(output_path)] if output_path.is_file() else [], "criterion_results": results, "reward": reward, "passed": passed}
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.write_text(f"{reward:.6f}\n", encoding="utf-8")
    raise SystemExit(0 if passed else 1)


if __name__ == "__main__":
    main()
