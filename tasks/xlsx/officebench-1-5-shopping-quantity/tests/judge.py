#!/usr/bin/env python3
"""确定性 Judge：从输入重算苹果数量并检查非目标保持。"""

from __future__ import annotations

import argparse
import hashlib
import json
import zipfile
from pathlib import Path

from openpyxl import load_workbook

TASK_ID = "cowork-evo/officebench-1-5-shopping-quantity"
RUBRIC_VERSION = "officebench-1-5-shopping-quantity-v0.1"
JUDGE_VERSION = "officebench-1-5-shopping-quantity-deterministic-v0.1"
INPUT_NAME = "shopping_list_input.xlsx"
OUTPUT_NAME = "shopping_list.xlsx"
EXPECTED_INPUT_SHA256 = "e6468e5452830bc4bd8767fb9a95c0ea95a65d876e306581437889e43c2a821a"
WEIGHTS = {"R001": 0.20, "R002": 0.35, "R003": 0.20, "R004": 0.10, "R005": 0.10, "R006": 0.05}


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def criterion(cid: str, score: bool | float, evidence: str) -> dict:
    return {"criterion_id": cid, "score": float(score), "evidence": evidence, "method": "deterministic"}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    input_path = Path(__file__).parents[1] / "data/input_files" / INPUT_NAME
    output_path = args.artifact_dir / OUTPUT_NAME
    results: list[dict] = []
    try:
        if not output_path.is_file() or not zipfile.is_zipfile(output_path):
            raise ValueError("shopping_list.xlsx 缺失或不是有效 OOXML")
        source = load_workbook(input_path, data_only=False)
        candidate = load_workbook(output_path, data_only=False)
        sws, cws = source["Sheet1"], candidate["Sheet1"]
        structure_ok = candidate.sheetnames == source.sheetnames and cws.max_row == sws.max_row and cws.max_column == sws.max_column
        results.append(criterion("R001", structure_ok, f"工作表={candidate.sheetnames}; 尺寸={cws.max_row}x{cws.max_column}"))
        target_ok = cws["A2"].value == "apple" and cws["B2"].value == sws["B2"].value / 2
        results.append(criterion("R002", target_ok, f"A2={cws['A2'].value!r}; B2={cws['B2'].value!r}; 期望={sws['B2'].value / 2}"))
        changed = []
        for row in range(1, sws.max_row + 1):
            for col in range(1, sws.max_column + 1):
                coord = sws.cell(row, col).coordinate
                if coord == "B2":
                    continue
                if sws.cell(row, col).value != cws.cell(row, col).value:
                    changed.append(coord)
        results.append(criterion("R003", not changed, f"非目标变更={changed[:10]}"))
        style_ok = all(sws.cell(row, col)._style == cws.cell(row, col)._style for row in range(1, sws.max_row + 1) for col in range(1, sws.max_column + 1) if sws.cell(row, col).coordinate != "B2")
        results.append(criterion("R004", style_ok, f"非目标样式保持={style_ok}"))
        results.append(criterion("R005", digest(input_path) == EXPECTED_INPUT_SHA256, f"输入 SHA256={digest(input_path)}"))
        results.append(criterion("R006", cws.protection.sheet is not True, "工作表可编辑"))
        source.close(); candidate.close()
    except Exception as exc:
        results = [criterion("R001", False, f"Judge 读取失败：{exc}")]
    by_id = {item["criterion_id"]: item["score"] for item in results}
    reward = sum(WEIGHTS[cid] * by_id.get(cid, 0.0) for cid in WEIGHTS)
    passed = reward >= 0.7 and all(by_id.get(cid) == 1.0 for cid in ("R001", "R002"))
    payload = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [digest(output_path)] if output_path.is_file() else [], "criterion_results": results, "reward": reward, "passed": passed}
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.write_text(f"{reward:.6f}\n", encoding="utf-8")
    raise SystemExit(0 if passed else 1)


if __name__ == "__main__":
    main()
