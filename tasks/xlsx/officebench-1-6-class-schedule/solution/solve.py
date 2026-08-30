#!/usr/bin/env python3
"""OfficeBench 1-6/2 的逐题参考解法。"""

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

from openpyxl import load_workbook

INPUT_NAME = "class_member.xlsx"
OUTPUT_NAME = "class_member_schedule.xlsx"
EXPECTED_INPUT_SHA256 = "7d62695b7c450cbe61059bf7e32e7dd0f901718e8f5de1e6fd34714123e2e85e"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    source = args.input_dir / INPUT_NAME
    if sha256(source) != EXPECTED_INPUT_SHA256:
        raise ValueError("class_member.xlsx 输入哈希不匹配")
    workbook = load_workbook(source)
    if workbook.sheetnames != ["Sheet1"]:
        raise ValueError(f"未知工作表结构：{workbook.sheetnames}")
    sheet = workbook["Sheet1"]
    if sheet.max_row != 6 or sheet["A1"].value != "Name":
        raise ValueError("输入表不是预期的 5 行姓名表")
    sheet["B1"] = "Meeting"
    for row in range(2, sheet.max_row + 1):
        minutes = (row - 2) * 10
        sheet.cell(row, 2).value = f"{8 + minutes // 60}:{minutes % 60:02d}" if minutes else "8:00"
        sheet.cell(row, 2).number_format = "@"
    sheet.column_dimensions["B"].width = 14
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    args.output_dir.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output_dir / OUTPUT_NAME)
    workbook.close()


if __name__ == "__main__":
    main()
