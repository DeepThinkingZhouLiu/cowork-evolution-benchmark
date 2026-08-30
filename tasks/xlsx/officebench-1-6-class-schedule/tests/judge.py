#!/usr/bin/env python3
"""确定性 Judge：检查排诊列的顺序、时间间隔和输入保持。"""

from __future__ import annotations

import argparse
import hashlib
import json
import zipfile
from pathlib import Path

from openpyxl import load_workbook

TASK_ID = "cowork-evo/officebench-1-6-class-schedule"
RUBRIC_VERSION = "officebench-1-6-v0.1"
JUDGE_VERSION = "officebench-1-6-deterministic-v0.1"
OUTPUT_NAME = "class_member_schedule.xlsx"
INPUT_NAME = "class_member.xlsx"
EXPECTED_INPUT_SHA256 = "7d62695b7c450cbe61059bf7e32e7dd0f901718e8f5de1e6fd34714123e2e85e"


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def c(cid, score, evidence):
    return {"criterion_id": cid, "score": float(score), "evidence": evidence, "method": "deterministic"}


def main():
    parser = argparse.ArgumentParser(); parser.add_argument("--artifact-dir", type=Path, required=True); parser.add_argument("--result", type=Path, required=True); parser.add_argument("--reward-file", type=Path, required=True); args = parser.parse_args()
    source_path = Path(__file__).parents[1] / "data/input_files" / INPUT_NAME
    out = args.artifact_dir / OUTPUT_NAME
    if not out.is_file():
        criteria = [c("R001", 0, "输出文件缺失")]; reward = 0.0
    else:
        try:
            if not zipfile.is_zipfile(out): raise ValueError("不是 OOXML ZIP")
            source = load_workbook(source_path, data_only=False); candidate = load_workbook(out, data_only=False)
            sw, cw = source["Sheet1"], candidate["Sheet1"]
            structure_ok = candidate.sheetnames == ["Sheet1"] and cw.max_row == 6 and cw.max_column >= 2
            criteria = [c("R001", structure_ok, f"工作表={candidate.sheetnames}, 尺寸={cw.max_row}x{cw.max_column}")]
            expected = ["8:00", "8:10", "8:20", "8:30", "8:40"]
            meeting = [cw.cell(row, 2).value for row in range(2, 7)]
            criteria.append(c("R002", cw["B1"].value == "Meeting" and meeting == expected, f"B1={cw['B1'].value!r}; Meeting={meeting!r}"))
            names_ok = [sw.cell(row, 1).value for row in range(1, 7)] == [cw.cell(row, 1).value for row in range(1, 7)]
            criteria.append(c("R003", names_ok, f"姓名保持={names_ok}"))
            changed = []
            for row in range(1, 7):
                for col in range(1, 2):
                    if sw.cell(row, col).value != cw.cell(row, col).value: changed.append(sw.cell(row, col).coordinate)
            criteria.append(c("R004", not changed, f"A 列非目标变更={changed}"))
            criteria.append(c("R005", digest(source_path) == EXPECTED_INPUT_SHA256, f"输入 SHA256={digest(source_path)}"))
            criteria.append(c("R006", cw.protection.sheet is not True and cw["B1"].value is not None, "工作表可编辑"))
            source.close(); candidate.close(); reward = sum(x["score"] for x in criteria) / len(criteria)
        except Exception as exc:
            criteria = [c("R001", 0, f"读取失败：{exc}")]; reward = 0.0
    passed = reward >= 0.7 and len(criteria) >= 2 and criteria[1]["score"] == 1.0
    payload = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [digest(out)] if out.is_file() else [], "criterion_results": criteria, "reward": reward, "passed": passed}
    args.result.parent.mkdir(parents=True, exist_ok=True); args.result.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.parent.mkdir(parents=True, exist_ok=True); args.reward_file.write_text(f"{reward:.6f}\n", encoding="utf-8")
    raise SystemExit(0 if passed else 1)


if __name__ == "__main__": main()
