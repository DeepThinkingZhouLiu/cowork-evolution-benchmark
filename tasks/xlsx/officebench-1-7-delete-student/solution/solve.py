#!/usr/bin/env python3
"""OfficeBench 1-7/0 的逐题参考解法。"""

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

from openpyxl import load_workbook

INPUT_NAME = "score_input.xlsx"
OUTPUT_NAME = "score.xlsx"
EXPECTED_INPUT_SHA256 = "023d650c0ce6192318d49f53cf099fddb109eb026953c96f3372e66fcbe01bc7"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    source = args.input_dir / INPUT_NAME
    if sha256(source) != EXPECTED_INPUT_SHA256:
        raise ValueError("score.xlsx 输入哈希不匹配")
    workbook = load_workbook(source)
    sheet = workbook["Sheet1"]
    delete_row = next((row for row in range(2, sheet.max_row + 1) if sheet.cell(row, 1).value == "Alice"), None)
    if delete_row is None:
        raise ValueError("输入中未找到 Alice")
    sheet.delete_rows(delete_row, 1)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    args.output_dir.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output_dir / OUTPUT_NAME)
    workbook.close()


if __name__ == "__main__":
    main()
