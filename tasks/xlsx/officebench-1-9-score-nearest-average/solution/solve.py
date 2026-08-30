#!/usr/bin/env python3
"""OfficeBench 1-9/1 的逐题参考解法。"""

from __future__ import annotations

import argparse
import hashlib
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

INPUT_NAME = "score_input.xlsx"
OUTPUT_NAME = "score.xlsx"
EXPECTED_INPUT_SHA256 = "12fce0b4e70006d9dc40745b532d75ee2e3c0ffc638cca2393bf2a9d5770226d"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    source = args.input_dir / INPUT_NAME
    if sha256(source) != EXPECTED_INPUT_SHA256:
        raise ValueError("score.xlsx 输入哈希不匹配")
    workbook = load_workbook(source)
    sheet = workbook["Sheet1"]
    records = [[sheet.cell(row, col).value for col in range(1, 4)] for row in range(2, sheet.max_row + 1)]
    average = sum(float(record[1]) for record in records) / len(records)
    nearest = min(enumerate(records), key=lambda item: (abs(float(item[1][1]) - average), item[0]))[1]
    append_row = sheet.max_row + 1
    for col, value in enumerate(nearest, start=1):
        target = sheet.cell(append_row, col, value)
        source_cell = sheet.cell(2, col)
        target.font = copy(source_cell.font)
        target.fill = copy(source_cell.fill)
        target.border = copy(source_cell.border)
        target.alignment = copy(source_cell.alignment)
        target.protection = copy(source_cell.protection)
        target.number_format = source_cell.number_format
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    args.output_dir.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output_dir / OUTPUT_NAME)
    workbook.close()


if __name__ == "__main__":
    main()
