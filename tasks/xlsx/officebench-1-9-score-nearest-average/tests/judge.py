#!/usr/bin/env python3
"""确定性 Judge：从输入计算平均值和最接近记录。"""

from __future__ import annotations

import argparse
import hashlib
import json
import zipfile
from pathlib import Path

from openpyxl import load_workbook

TASK_ID = "cowork-evo/officebench-1-9-score-nearest-average"
RUBRIC_VERSION = "officebench-1-9-score-nearest-average-v0.1"
JUDGE_VERSION = "officebench-1-9-score-nearest-average-deterministic-v0.1"
INPUT_NAME = "score_input.xlsx"
OUTPUT_NAME = "score.xlsx"
EXPECTED_INPUT_SHA256 = "12fce0b4e70006d9dc40745b532d75ee2e3c0ffc638cca2393bf2a9d5770226d"
WEIGHTS = {"R001": 0.20, "R002": 0.35, "R003": 0.20, "R004": 0.15, "R005": 0.05, "R006": 0.05}


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def c(cid: str, score: bool | float, evidence: str) -> dict:
    return {"criterion_id": cid, "score": float(score), "evidence": evidence, "method": "deterministic"}


def main() -> None:
    parser = argparse.ArgumentParser(); parser.add_argument("--artifact-dir", type=Path, required=True); parser.add_argument("--result", type=Path, required=True); parser.add_argument("--reward-file", type=Path, required=True); args = parser.parse_args()
    input_path = Path(__file__).parents[1] / "data/input_files" / INPUT_NAME
    output_path = args.artifact_dir / OUTPUT_NAME
    results: list[dict] = []
    try:
        if not output_path.is_file() or not zipfile.is_zipfile(output_path):
            raise ValueError("score.xlsx 缺失或不是有效 OOXML")
        source = load_workbook(input_path, data_only=False); candidate = load_workbook(output_path, data_only=False)
        sws, cws = source["Sheet1"], candidate["Sheet1"]
        source_rows = [[sws.cell(row, col).value for col in range(1, 4)] for row in range(2, sws.max_row + 1)]
        average = sum(float(row[1]) for row in source_rows) / len(source_rows)
        nearest_index, nearest = min(enumerate(source_rows), key=lambda item: (abs(float(item[1][1]) - average), item[0]))
        expected_rows = source_rows + [nearest]
        actual_rows = [[cws.cell(row, col).value for col in range(1, 4)] for row in range(2, cws.max_row + 1)]
        structure_ok = candidate.sheetnames == source.sheetnames and cws.max_row == sws.max_row + 1 and cws.max_column == sws.max_column
        results.append(c("R001", structure_ok, f"工作表={candidate.sheetnames}; 尺寸={cws.max_row}x{cws.max_column}"))
        logical_ok = actual_rows == expected_rows
        results.append(c("R002", logical_ok, f"midterm1 平均值={average:.2f}; 最接近输入行={nearest_index + 2}; 追加矩阵一致={logical_ok}"))
        record_ok = actual_rows[-1] == nearest
        results.append(c("R003", record_ok, f"追加记录={actual_rows[-1] if actual_rows else None}; 期望={nearest}"))
        original_values_ok = actual_rows[:-1] == source_rows
        original_style_ok = all(sws.cell(row, col)._style == cws.cell(row, col)._style for row in range(1, sws.max_row + 1) for col in range(1, sws.max_column + 1))
        results.append(c("R004", original_values_ok and original_style_ok, f"原有值保持={original_values_ok}; 原有样式保持={original_style_ok}"))
        results.append(c("R005", digest(input_path) == EXPECTED_INPUT_SHA256, f"输入 SHA256={digest(input_path)}"))
        results.append(c("R006", cws.protection.sheet is not True, "工作表可编辑"))
        source.close(); candidate.close()
    except Exception as exc:
        results = [c("R001", False, f"Judge 读取失败：{exc}")]
    by_id = {item["criterion_id"]: item["score"] for item in results}; reward = sum(WEIGHTS[key] * by_id.get(key, 0.0) for key in WEIGHTS)
    passed = reward >= 0.7 and all(by_id.get(key) == 1.0 for key in ("R001", "R002", "R003"))
    payload = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [digest(output_path)] if output_path.is_file() else [], "criterion_results": results, "reward": reward, "passed": passed}
    args.result.parent.mkdir(parents=True, exist_ok=True); args.result.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.parent.mkdir(parents=True, exist_ok=True); args.reward_file.write_text(f"{reward:.6f}\n", encoding="utf-8")
    raise SystemExit(0 if passed else 1)


if __name__ == "__main__": main()
