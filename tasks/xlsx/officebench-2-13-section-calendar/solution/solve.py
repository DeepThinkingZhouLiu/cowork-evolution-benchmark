#!/usr/bin/env python3
"""OfficeBench 2-13/1 的逐题参考解法。"""

from __future__ import annotations

import argparse
import hashlib
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

INPUT_NAME = "team_training_schedule.xlsx"
EXPECTED_INPUT_SHA256 = "a5628f45d9f48eccd8424634fda4d3c9048200571cc03db1e6d34cf9447498f4"
EVENT_DATE = "20200501"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def parse_time_range(value: str) -> tuple[str, str]:
    match = re.fullmatch(r"\s*(\d{1,2}):(\d{2})\s*AM\s*-\s*(\d{1,2}):(\d{2})\s*AM\s*", value)
    if not match:
        raise ValueError(f"无法解析时间范围：{value!r}")
    start_h, start_m, end_h, end_m = (int(x) for x in match.groups())
    return f"{start_h:02d}{start_m:02d}00", f"{end_h:02d}{end_m:02d}00"


def escape_ics(value: str) -> str:
    return value.replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")


def make_event(name: str, section: str, time_range: str) -> str:
    start, end = parse_time_range(time_range)
    uid = f"{name.lower()}-{EVENT_DATE}-section-a@coworkevobench.local"
    stamp = datetime.utcnow().strftime("%Y%m%dT%H%M%SZ")
    lines = [
        "BEGIN:VCALENDAR",
        "VERSION:2.0",
        "PRODID:-//CoworkEvoBench//OfficeBench 2-13//EN",
        "CALSCALE:GREGORIAN",
        "METHOD:PUBLISH",
        "BEGIN:VEVENT",
        f"UID:{escape_ics(uid)}",
        f"DTSTAMP:{stamp}",
        f"DTSTART:{EVENT_DATE}T{start}",
        f"DTEND:{EVENT_DATE}T{end}",
        f"SUMMARY:{escape_ics('Team training session - ' + name)}",
        f"DESCRIPTION:{escape_ics('Participant: ' + name + '; Section: ' + section)}",
        "END:VEVENT",
        "END:VCALENDAR",
        "",
    ]
    return "\r\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    source = args.input_dir / INPUT_NAME
    if sha256(source) != EXPECTED_INPUT_SHA256:
        raise ValueError("team_training_schedule.xlsx 输入哈希不匹配")
    workbook = load_workbook(source, data_only=True, read_only=True)
    if workbook.sheetnames != ["Training Schedule"]:
        raise ValueError(f"未知工作表结构：{workbook.sheetnames}")
    sheet = workbook["Training Schedule"]
    headers = [sheet.cell(1, col).value for col in range(1, 4)]
    if headers != ["Name", "Section", "Time"]:
        raise ValueError(f"未知表头：{headers!r}")
    rows = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        if row[1] == "Section A":
            rows.append((str(row[0]), str(row[1]), str(row[2])))
    workbook.close()
    if [row[0] for row in rows] != ["Alice", "David", "Grace", "Jack"]:
        raise ValueError(f"Section A 成员与固定快照不符：{rows!r}")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for name, section, time_range in rows:
        (args.output_dir / f"{name}.ics").write_text(make_event(name, section, time_range), encoding="utf-8", newline="")


if __name__ == "__main__":
    main()
