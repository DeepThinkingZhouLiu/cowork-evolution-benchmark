#!/usr/bin/env python3
"""从输入工作簿独立重算 Section A 事件并检查 iCalendar 语义。"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

from openpyxl import load_workbook

TASK_ID = "cowork-evo/officebench-2-13-section-calendar"
RUBRIC_VERSION = "officebench-2-13-v0.1"
JUDGE_VERSION = "officebench-2-13-deterministic-v0.1"
INPUT_NAME = "team_training_schedule.xlsx"
EXPECTED_INPUT_SHA256 = "a5628f45d9f48eccd8424634fda4d3c9048200571cc03db1e6d34cf9447498f4"
EXPECTED_FILES = ["Alice.ics", "David.ics", "Grace.ics", "Jack.ics"]


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def criterion(cid: str, score: float, evidence: str) -> dict:
    return {"criterion_id": cid, "score": float(score), "evidence": evidence, "method": "deterministic"}


def parse_event(path: Path) -> dict[str, str]:
    text = path.read_bytes().decode("utf-8")
    if not text.endswith("\r\n") or "BEGIN:VCALENDAR\r\n" not in text or "END:VCALENDAR\r\n" not in text:
        raise ValueError("缺少完整 VCALENDAR 边界")
    if text.count("BEGIN:VEVENT\r\n") != 1 or text.count("END:VEVENT\r\n") != 1:
        raise ValueError("VEVENT 数量不是 1")
    event = text.split("BEGIN:VEVENT\r\n", 1)[1].split("\r\nEND:VEVENT\r\n", 1)[0]
    props: dict[str, str] = {}
    for line in event.split("\r\n"):
        if ":" not in line:
            continue
        key, value = line.split(":", 1)
        props[key.split(";", 1)[0].upper()] = value.replace("\\,", ",").replace("\\;", ";").replace("\\n", "\n").replace("\\\\", "\\")
    required = {"UID", "DTSTAMP", "DTSTART", "DTEND", "SUMMARY", "DESCRIPTION"}
    if not required.issubset(props):
        raise ValueError(f"缺少属性：{sorted(required - props)}")
    return props


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    source_path = Path(__file__).parents[1] / "data/input_files" / INPUT_NAME
    results: list[dict] = []
    try:
        source = load_workbook(source_path, data_only=True, read_only=True)
        sheet = source["Training Schedule"]
        rows = [row for row in sheet.iter_rows(min_row=2, values_only=True) if row[0] is not None and row[1] == "Section A"]
        expected = {str(row[0]): (str(row[2]), "20200501") for row in rows}
        source.close()
        expected_names = ["Alice", "David", "Grace", "Jack"]
        expected_ok = list(expected) == expected_names and all(value[0] == "9:00 AM - 11:00 AM" for value in expected.values())
        names = sorted(p.stem for p in args.artifact_dir.glob("*.ics"))
        missing = [name for name in EXPECTED_FILES if not (args.artifact_dir / name).is_file()]
        extras = [name for name in names if f"{name}.ics" not in EXPECTED_FILES]
        results.append(criterion("R001", 1.0 if not missing else 0.0, f"缺失文件={missing}"))
        results.append(criterion("R002", 1.0 if expected_ok and names == sorted(name[:-4] for name in EXPECTED_FILES) and not extras else 0.0, f"输入成员={list(expected)}; 输出 stems={names}; 额外={extras}"))
        parsed: dict[str, dict[str, str]] = {}
        parse_errors = []
        for filename in EXPECTED_FILES:
            path = args.artifact_dir / filename
            if not path.is_file():
                continue
            try:
                parsed[filename[:-4]] = parse_event(path)
            except Exception as exc:
                parse_errors.append(f"{filename}: {exc}")
        timing_ok = all(props.get("DTSTART") == "20200501T090000" and props.get("DTEND") == "20200501T110000" for props in parsed.values()) and len(parsed) == 4
        results.append(criterion("R003", 1.0 if timing_ok and not parse_errors else 0.0, f"时间检查={timing_ok}; 解析错误={parse_errors}"))
        content_ok = all(name in props.get("SUMMARY", "") and "Team training session" in props.get("SUMMARY", "") and name in props.get("DESCRIPTION", "") and "Section A" in props.get("DESCRIPTION", "") for name, props in parsed.items()) and len(parsed) == 4
        results.append(criterion("R004", 1.0 if content_ok else 0.0, f"标题/描述检查={content_ok}"))
        native_ok = all("BEGIN:VEVENT" in (args.artifact_dir / f"{name}.ics").read_text(encoding="utf-8") for name in parsed)
        results.append(criterion("R005", 1.0 if native_ok and not parse_errors else 0.0, f"原生 VEVENT={native_ok}"))
        input_hash_ok = digest(source_path) == EXPECTED_INPUT_SHA256
        results.append(criterion("R006", 1.0 if input_hash_ok else 0.0, f"输入 SHA256={digest(source_path)}"))
        reward = sum(item["score"] for item in results) / len(results)
    except Exception as exc:
        results = [criterion("R001", 0.0, f"Judge 读取失败：{exc}")]
        reward = 0.0
    passed = reward >= 0.7 and len(results) >= 3 and results[0]["score"] == 1.0 and results[1]["score"] == 1.0 and results[2]["score"] == 1.0
    payload = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [digest(path) for path in sorted(args.artifact_dir.glob("*.ics"))], "criterion_results": results, "reward": reward, "passed": passed}
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.write_text(f"{reward:.6f}\n", encoding="utf-8")
    raise SystemExit(0 if passed else 1)


if __name__ == "__main__":
    main()
