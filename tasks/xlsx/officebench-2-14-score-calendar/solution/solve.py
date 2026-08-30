#!/usr/bin/env python3
"""从成绩工作簿独立计算前三名并生成三个原生 iCalendar 事件。"""

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

from openpyxl import load_workbook

INPUT_NAME = "score_input.xlsx"
EXPECTED_SHA256 = "a5c8026dc09f38411c6b8783dad74f7f66f08387eeb21b9d8eb7f28a90dd5b6e"
DATE = "20200501"
TIME = "100000"
DTSTAMP = "20200101T000000Z"


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def escape_ics(value: str) -> str:
    return value.replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")


def event(name: str, rank: int) -> str:
    uid = f"{name.lower()}-{DATE}-reward@coworkevobench.local"
    lines = [
        "BEGIN:VCALENDAR",
        "VERSION:2.0",
        "PRODID:-//CoworkEvoBench//OfficeBench 2-14//EN",
        "CALSCALE:GREGORIAN",
        "METHOD:PUBLISH",
        "BEGIN:VEVENT",
        f"UID:{uid}",
        f"DTSTAMP:{DTSTAMP}",
        f"DTSTART:{DATE}T{TIME}",
        f"SUMMARY:{escape_ics('Reward meeting - ' + name)}",
        f"DESCRIPTION:{escape_ics(f'Reward meeting for {name}, ranked top-three by midterm1 (rank {rank}).')}",
        "END:VEVENT",
        "END:VCALENDAR",
        "",
    ]
    return "\r\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    source_path = args.input_dir / INPUT_NAME
    if digest(source_path) != EXPECTED_SHA256:
        raise ValueError("score_input.xlsx 输入哈希不匹配")

    workbook = load_workbook(source_path, data_only=True, read_only=True)
    if workbook.sheetnames != ["Sheet1"]:
        raise ValueError(f"工作表结构不符合输入快照：{workbook.sheetnames}")
    sheet = workbook["Sheet1"]
    headers = [sheet.cell(1, col).value for col in range(1, 4)]
    if headers != ["Name", "midterm1", "midterm2"]:
        raise ValueError(f"输入表头不符合预期：{headers!r}")
    rows = []
    for row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if row[0] is not None:
            rows.append((row_index, str(row[0]), float(row[1])))
    workbook.close()
    top = sorted(rows, key=lambda item: (-item[2], item[0]))[:3]
    if [(name, score) for _, name, score in top] != [("Tina", 100.0), ("Rachel", 100.0), ("Sam", 99.0)]:
        raise ValueError(f"前三名与固定输入不符：{top!r}")

    args.output_dir.mkdir(parents=True, exist_ok=True)
    for rank, (_, name, _) in enumerate(top, start=1):
        (args.output_dir / f"{name}.ics").write_text(event(name, rank), encoding="utf-8", newline="")


if __name__ == "__main__":
    main()
