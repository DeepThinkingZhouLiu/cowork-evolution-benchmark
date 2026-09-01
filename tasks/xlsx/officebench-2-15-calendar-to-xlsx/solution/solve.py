#!/usr/bin/env python3
"""Reference solution for OfficeBench 2-15/0."""
from __future__ import annotations
import argparse, hashlib, re
from datetime import datetime, time
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

INPUT = "Alice.ics"
EXPECTED_SHA256 = "a9d16392b8a24ad049655586076f0fa657890af5f3b1c85e0139d602514b2f13"

def digest(path: Path) -> str:
    h = hashlib.sha256(); h.update(path.read_bytes()); return h.hexdigest()

def parse(path: Path):
    text = path.read_text(encoding="utf-8")
    events = []
    for block in re.findall(r"BEGIN:VEVENT\r?\n(.*?)END:VEVENT", text, flags=re.S):
        props = {}
        for line in block.splitlines():
            if ":" in line:
                key, value = line.split(":", 1); props[key.split(";", 1)[0].upper()] = value
        start = datetime.strptime(props["DTSTART"].strip(), "%Y%m%dT%H%M%SZ")
        end = datetime.strptime(props["DTEND"].strip(), "%Y%m%dT%H%M%SZ")
        summary = props["SUMMARY"]
        item, contact = summary.rsplit(" by ", 1)
        events.append((start, end, item, contact))
    return sorted(events)

def main() -> None:
    parser = argparse.ArgumentParser(); parser.add_argument("--input-dir", type=Path, required=True); parser.add_argument("--output-dir", type=Path, required=True); args = parser.parse_args()
    source = args.input_dir / INPUT
    if digest(source) != EXPECTED_SHA256: raise ValueError("calendar input hash mismatch")
    events = parse(source)
    if len(events) != 3: raise ValueError("expected three calendar events")
    wb = Workbook(); ws = wb.active; ws.title = "Meeting agenda"
    ws.merge_cells("B1:F1"); ws["B1"] = "MEETING AGENDA"; ws["B1"].font = Font(bold=True, size=14); ws["B1"].alignment = Alignment(horizontal="center")
    headers = ["Start", "End", "Time", "Item", "Contact"]
    for col, header in enumerate(headers, 2):
        cell = ws.cell(2, col, header); cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row, (start, end, item, contact) in enumerate(events, 3):
        ws.cell(row, 2, start.time()); ws.cell(row, 3, end.time()); ws.cell(row, 4, (end-start).seconds / 86400); ws.cell(row, 5, item); ws.cell(row, 6, contact)
        for col in (2, 3): ws.cell(row, col).number_format = "h:mm AM/PM"
        ws.cell(row, 4).number_format = "h:mm"
    ws.freeze_panes = "B3"; ws.auto_filter.ref = "B2:F5"; ws.print_area = "B1:F5"; ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
    for col, width in {"B":14, "C":14, "D":10, "E":20, "F":16}.items(): ws.column_dimensions[col].width = width
    args.output_dir.mkdir(parents=True, exist_ok=True); wb.save(args.output_dir / "meeting_agenda.xlsx")

if __name__ == "__main__": main()
