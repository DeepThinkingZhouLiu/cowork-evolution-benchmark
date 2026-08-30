#!/usr/bin/env python3
"""Independent semantic Judge for calendar-to-workbook extraction."""
from __future__ import annotations
import argparse, hashlib, json, re, zipfile
from datetime import datetime, time
from pathlib import Path
from openpyxl import load_workbook

TASK_ID = "cowork-evo/officebench-2-15-calendar-to-xlsx"; RUBRIC_VERSION = "officebench-2-15-calendar-to-xlsx-v0.1"; JUDGE_VERSION = "officebench-2-15-calendar-to-xlsx-deterministic-v0.1"
EXPECTED_SHA256 = "a9d16392b8a24ad049655586076f0fa657890af5f3b1c85e0139d602514b2f13"; WEIGHTS = {"R001": .15, "R002": .15, "R003": .25, "R004": .25, "R005": .10, "R006": .10}

def digest(path: Path) -> str: return hashlib.sha256(path.read_bytes()).hexdigest()
def crit(cid, score, evidence): return {"criterion_id": cid, "score": float(score), "evidence": evidence, "method": "deterministic"}
def seconds(value):
    if isinstance(value, datetime): return value.hour*3600+value.minute*60+value.second
    if isinstance(value, time): return value.hour*3600+value.minute*60+value.second
    if isinstance(value, (int,float)): return round(float(value)*86400)
    return None

def source_events(path):
    text = path.read_text(encoding="utf-8"); events=[]
    for block in re.findall(r"BEGIN:VEVENT\r?\n(.*?)END:VEVENT", text, flags=re.S):
        props={}
        for line in block.splitlines():
            if ":" in line: key,val=line.split(":",1); props[key.split(";",1)[0].upper()]=val
        start=datetime.strptime(props["DTSTART"], "%Y%m%dT%H%M%SZ"); end=datetime.strptime(props["DTEND"], "%Y%m%dT%H%M%SZ"); item,contact=props["SUMMARY"].rsplit(" by ",1); events.append((start,end,item,contact))
    return sorted(events)

def main():
    ap=argparse.ArgumentParser(); ap.add_argument("--artifact-dir",type=Path,required=True); ap.add_argument("--result",type=Path,required=True); ap.add_argument("--reward-file",type=Path,required=True); a=ap.parse_args(); results=[]
    source=Path(__file__).parents[1]/"data/input_files/Alice.ics"; out=a.artifact_dir/"meeting_agenda.xlsx"
    try:
        missing=[] if out.is_file() else ["meeting_agenda.xlsx"]; native=bool(out.is_file() and zipfile.is_zipfile(out)); results.append(crit("R001", native, f"missing={missing}; OOXML={native}"))
        expected=source_events(source); wb=load_workbook(out,data_only=True) if native else None
        ws=wb["Meeting agenda"] if wb and "Meeting agenda" in wb.sheetnames else (wb.active if wb else None)
        headers=tuple(ws.cell(2,c).value for c in range(2,7)) if ws else (); structure=bool(ws and headers==("Start","End","Time","Item","Contact")); results.append(crit("R002",structure,f"sheet={wb.sheetnames if wb else []}; headers={headers}"))
        rows=[]
        if ws:
            for r in range(3, ws.max_row+1):
                vals=[ws.cell(r,c).value for c in range(2,7)]
                if any(v is not None for v in vals): rows.append(vals)
        expected_rows=[(e[0].time(),e[1].time(),(e[1]-e[0]).seconds/86400,e[2],e[3]) for e in expected]
        row_shape=len(rows)==3 and [seconds(x[0]) for x in rows]==[seconds(x[0]) for x in expected_rows]; results.append(crit("R003",row_shape,f"rows={len(rows)}; ordered={row_shape}"))
        fidelity=len(rows)==3 and all(seconds(g[0])==seconds(e[0]) and seconds(g[1])==seconds(e[1]) and abs((float(g[2]) if isinstance(g[2],(int,float)) else seconds(g[2])/86400)-float(e[2]))<1e-6 and str(g[3])==e[3] and str(g[4])==e[4] for g,e in zip(rows,expected_rows)); results.append(crit("R004",fidelity,f"expected={expected_rows}; output={rows}"))
        formatted=bool(ws and ws.merged_cells.ranges and str(ws["B1"].value).upper()=="MEETING AGENDA" and all(ws.cell(r,c).number_format != "General" for r in range(3,6) for c in (2,3,4))); results.append(crit("R005",formatted,f"merged={list(ws.merged_cells.ranges) if ws else []}"))
        hash_ok=digest(source)==EXPECTED_SHA256; results.append(crit("R006",hash_ok,f"input_sha256={digest(source)}"))
        if wb: wb.close()
    except Exception as exc: results=[crit("R001",False,f"Judge failure: {exc}")]
    values={x["criterion_id"]:x["score"] for x in results}; reward=sum(WEIGHTS[k]*values.get(k,0) for k in WEIGHTS); passed=reward>=.8 and all(values.get(k)==1.0 for k in ("R001","R003","R004"))
    payload={"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":[digest(out)] if out.is_file() else [],"criterion_results":results,"reward":reward,"passed":passed}; a.result.parent.mkdir(parents=True,exist_ok=True); a.result.write_text(json.dumps(payload,indent=2),encoding="utf-8"); a.reward_file.parent.mkdir(parents=True,exist_ok=True); a.reward_file.write_text(f"{reward:.6f}\n"); raise SystemExit(0 if passed else 1)

if __name__ == "__main__": main()
