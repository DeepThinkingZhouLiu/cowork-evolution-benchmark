#!/usr/bin/env python3
from __future__ import annotations
import argparse
from collections import defaultdict
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import load_workbook

INPUT_NAME = "student_scores.xlsx"

def read_rows(path: Path):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["Sheet1"]
    rows = [(str(name), float(score)) for name, score in ws.iter_rows(min_row=2, max_col=2, values_only=True) if name is not None]
    wb.close()
    return rows

def band(score: float) -> str:
    return "A" if score >= 90 else "B" if score >= 80 else "C" if score >= 70 else "D" if score >= 60 else "F"

def solve(input_dir: Path, output_dir: Path) -> None:
    rows = read_rows(input_dir / INPUT_NAME)
    groups = defaultdict(list)
    for name, score in rows: groups[name].append(score)
    output_dir.mkdir(parents=True, exist_ok=True)
    doc = Document()
    title = doc.add_heading("Student Grade Teaching Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Prepared from student_scores.xlsx; repeated names are consolidated by student.")
    doc.add_heading("Class Overview", level=1)
    dist = {key: sum(1 for _, score in rows if band(score) == key) for key in ["A", "B", "C", "D", "F"]}
    doc.add_paragraph(f"Grade records: {len(rows)} | Students: {len(groups)} | Class average: {sum(s for _, s in rows) / len(rows):.2f}")
    table = doc.add_table(rows=1, cols=2); table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.rows[0].cells[0].text = "Band"; table.rows[0].cells[1].text = "Records"
    for key in ["A", "B", "C", "D", "F"]:
        cells = table.add_row().cells; cells[0].text = key; cells[1].text = str(dist[key])
    doc.add_heading("Student Summary", level=1)
    table = doc.add_table(rows=1, cols=5); table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, ["Student", "Records", "Average", "Highest", "Lowest"]): cell.text = text
    for name in sorted(groups):
        scores = groups[name]; cells = table.add_row().cells
        cells[0].text = name; cells[1].text = str(len(scores)); cells[2].text = f"{sum(scores)/len(scores):.2f}"; cells[3].text = f"{max(scores):.0f}"; cells[4].text = f"{min(scores):.0f}"
    doc.add_heading("Teaching Priorities", level=1)
    support = sorted((name, sum(scores)/len(scores)) for name, scores in groups.items() if sum(scores)/len(scores) < 70)
    praise = sorted((name, sum(scores)/len(scores)) for name, scores in groups.items() if sum(scores)/len(scores) >= 90)
    doc.add_paragraph("Priority support: " + (", ".join(f"{n} ({s:.2f})" for n, s in support) if support else "No student below 70 average.") + ".")
    doc.add_paragraph("High-performing students: " + (", ".join(f"{n} ({s:.2f})" for n, s in praise) if praise else "No student at or above 90 average.") + ".")
    doc.add_paragraph("Use the lowest individual scores in the student summary to target review, and recognize sustained high performance.")
    doc.save(output_dir / "teaching.docx")

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--input-dir", type=Path, required=True); ap.add_argument("--output-dir", type=Path, required=True); a = ap.parse_args(); solve(a.input_dir, a.output_dir)
if __name__ == "__main__": main()
