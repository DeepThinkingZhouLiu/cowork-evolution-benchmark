#!/usr/bin/env python3
from __future__ import annotations
import argparse, hashlib, json, re, zipfile
from collections import Counter, defaultdict
from pathlib import Path
from docx import Document
from openpyxl import load_workbook

OUT = "teaching.docx"; TASK_ID = "cowork-evo/officebench-2-19-student-teaching-report"
WEIGHTS = {"R001":.15,"R002":.18,"R003":.22,"R004":.18,"R005":.12,"R006":.10,"R007":.05}
def crit(cid, ok, ev): return {"criterion_id":cid,"score":1.0 if ok else 0.0,"evidence":ev,"method":"deterministic"}
def band(s): return "A" if s>=90 else "B" if s>=80 else "C" if s>=70 else "D" if s>=60 else "F"
def main():
    ap=argparse.ArgumentParser(); ap.add_argument("--artifact-dir",type=Path,required=True); ap.add_argument("--result",type=Path,required=True); ap.add_argument("--reward-file",type=Path,required=True); a=ap.parse_args(); source=Path(__file__).parents[1]/"data/input_files/student_scores.xlsx"; out=a.artifact_dir/OUT; results=[]
    try:
        wb=load_workbook(source,data_only=True,read_only=True); ws=wb["Sheet1"]; rows=[(str(n),float(s)) for n,s in ws.iter_rows(min_row=2,max_col=2,values_only=True) if n is not None]; wb.close(); groups=defaultdict(list)
        for n,s in rows: groups[n].append(s)
        text=""; tables=[]
        if out.is_file() and zipfile.is_zipfile(out):
            doc=Document(out); text="\n".join([p.text for p in doc.paragraphs]);
            for t in doc.tables: tables.append([[c.text.strip() for c in row.cells] for row in t.rows]); text += "\n"+"\n".join(" ".join(x) for x in tables[-1])
        valid=out.is_file() and zipfile.is_zipfile(out) and bool(text)
        results.append(crit("R001",valid,"teaching.docx 原生 OOXML 且可解析"))
        expected_count=f"{len(rows)}"; avg=sum(s for _,s in rows)/len(rows)
        overview=expected_count in text and f"{avg:.2f}" in text and str(len(groups)) in text
        results.append(crit("R002",overview,f"记录={len(rows)}，学生={len(groups)}，平均={avg:.2f}"))
        dist=Counter(band(s) for _,s in rows); dist_ok=all(re.search(rf"\b{k}\b[^0-9]*{v}\b",text) for k,v in dist.items())
        results.append(crit("R003",bool(dist_ok),f"分布={dict(dist)}"))
        summary={r[0]:r[1:] for r in tables[-1][1:]} if tables and len(tables[-1][0])>=5 else {}
        student_ok=len(summary)==len(groups) and all(name in summary and int(float(summary[name][0]))==len(scores) and abs(float(summary[name][1])-sum(scores)/len(scores))<.011 and abs(float(summary[name][2])-max(scores))<.01 and abs(float(summary[name][3])-min(scores))<.01 for name,scores in groups.items())
        results.append(crit("R004",student_ok,"逐学生记录数、平均、最高和最低分独立重算"))
        support=[n for n,s in groups.items() if sum(s)/len(s)<70]; praise=[n for n,s in groups.items() if sum(s)/len(s)>=90]
        support_ok=("Priority support" in text and all(n in text for n in support)); praise_ok=("High-performing" in text and all(n in text for n in praise))
        results.append(crit("R005",support_ok and praise_ok,f"辅导={support}; 表扬={praise}"))
        results.append(crit("R006",hashlib.sha256(source.read_bytes()).hexdigest()=="e151e7aff861749d1fce7bea0e584263cdc5182df25f9b6a9a96cb62ca9b19f8","输入 SHA256 固定"))
        results.append(crit("R007",len(tables)>=2 and all(len(row)>=2 for row in tables[0]) and len(tables[-1])>=len(groups)+1,"至少包含概览和逐学生可编辑表格"))
    except Exception as exc: results=[crit("R001",False,f"Judge 读取失败：{exc}")]
    vals={x["criterion_id"]:x["score"] for x in results}; reward=sum(WEIGHTS[k]*vals.get(k,0) for k in WEIGHTS); passed=reward>=.8 and vals.get("R001",0)==1 and vals.get("R002",0)==1 and vals.get("R004",0)==1
    payload={"task_id":TASK_ID,"rubric_version":"officebench-2-19-v0.1","judge_version":"officebench-2-19-deterministic-v0.1","artifact_sha256":[hashlib.sha256(out.read_bytes()).hexdigest()] if out.is_file() else [],"criterion_results":results,"reward":reward,"passed":passed}; a.result.parent.mkdir(parents=True,exist_ok=True); a.result.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8"); a.reward_file.parent.mkdir(parents=True,exist_ok=True); a.reward_file.write_text(f"{reward:.6f}\n"); raise SystemExit(0 if passed else 1)
if __name__=="__main__": main()
