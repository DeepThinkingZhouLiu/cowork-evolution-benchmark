#!/usr/bin/env python3
"""Reference solution for OfficeBench 2-24/0."""
from __future__ import annotations
import argparse, hashlib
from copy import copy
from pathlib import Path
from openpyxl import load_workbook

INPUT = "score_input.xlsx"; EXPECTED_SHA256 = "e15533ba4ea8f4bcddf454c9e05d0f34fcc511a20041600c72b2a98ae9c3551d"
def digest(path: Path) -> str: return hashlib.sha256(path.read_bytes()).hexdigest()
def main() -> None:
    ap=argparse.ArgumentParser(); ap.add_argument("--input-dir",type=Path,required=True); ap.add_argument("--output-dir",type=Path,required=True); a=ap.parse_args(); source=a.input_dir/INPUT
    if digest(source)!=EXPECTED_SHA256: raise ValueError("score input hash mismatch")
    wb=load_workbook(source); ws=wb.active; rows=list(ws.iter_rows(min_row=2,max_col=2,values_only=True)); scores=[float(row[1]) for row in rows]
    ws.cell(ws.max_row+1,1,"Average"); ws.cell(ws.max_row,2,sum(scores)/len(scores)); ws.cell(ws.max_row,2).number_format="0.00"; font=copy(ws.cell(ws.max_row,2).font); font.bold=True; ws.cell(ws.max_row,1).font=font
    ws.auto_filter.ref=f"A1:B{ws.max_row}"; ws.freeze_panes="A2"; ws.column_dimensions["A"].width=16; ws.column_dimensions["B"].width=12
    a.output_dir.mkdir(parents=True,exist_ok=True); wb.save(a.output_dir/"score.xlsx"); wb.close()
if __name__=="__main__": main()
