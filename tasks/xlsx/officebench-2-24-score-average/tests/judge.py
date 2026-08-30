#!/usr/bin/env python3
"""Independent Judge for the score extraction and average."""
from __future__ import annotations
import argparse, hashlib, json, math, zipfile
from pathlib import Path
from openpyxl import load_workbook

TASK_ID="cowork-evo/officebench-2-24-score-average"; RUBRIC_VERSION="officebench-2-24-score-average-v0.1"; JUDGE_VERSION="officebench-2-24-score-average-deterministic-v0.1"; EXPECTED_SHA256="e15533ba4ea8f4bcddf454c9e05d0f34fcc511a20041600c72b2a98ae9c3551d"; WEIGHTS={"R001":.15,"R002":.25,"R003":.30,"R004":.15,"R005":.10,"R006":.05}
def digest(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def crit(cid,score,evidence): return {"criterion_id":cid,"score":float(score),"evidence":evidence,"method":"deterministic"}
def main():
    ap=argparse.ArgumentParser(); ap.add_argument("--artifact-dir",type=Path,required=True); ap.add_argument("--result",type=Path,required=True); ap.add_argument("--reward-file",type=Path,required=True); a=ap.parse_args(); out=a.artifact_dir/"score.xlsx"; src=Path(__file__).parents[1]/"data/input_files/score_input.xlsx"; results=[]
    try:
        native=out.is_file() and zipfile.is_zipfile(out); results.append(crit("R001",native,f"exists={out.is_file()}; OOXML={native}"))
        source_wb=load_workbook(src,data_only=True); source_ws=source_wb.active; source_rows=[(str(r[0]),float(r[1])) for r in source_ws.iter_rows(min_row=2,max_col=2,values_only=True)]; source_wb.close(); expected_avg=sum(x[1] for x in source_rows)/len(source_rows)
        wb=load_workbook(out,data_only=True) if native else None; ws=wb.active if wb else None; headers=tuple(ws.cell(1,c).value for c in (1,2)) if ws else (); rows=[]
        if ws:
            for r in range(2,ws.max_row+1):
                if ws.cell(r,1).value is not None or ws.cell(r,2).value is not None: rows.append((ws.cell(r,1).value,ws.cell(r,2).value))
        preserved=len(rows)>=4 and headers==("Name","Score") and [(str(n),float(s)) for n,s in rows[:3]]==source_rows; results.append(crit("R002",preserved,f"headers={headers}; source_rows={source_rows}; output_rows={rows}"))
        avg_ok=len(rows)==4 and str(rows[-1][0]).strip().lower()=="average" and isinstance(rows[-1][1],(int,float)) and math.isclose(float(rows[-1][1]),expected_avg,abs_tol=1e-9); results.append(crit("R003",avg_ok,f"expected_average={expected_avg}; output_average={rows[-1][1] if rows else None}"))
        no_dupes=len(rows)==4 and [(str(n),float(s)) for n,s in rows[:3]]==source_rows and all(str(n).strip().lower()!="average" for n,s in rows[:3]); results.append(crit("R004",no_dupes,f"row_count={len(rows)}"))
        professional=bool(ws and ws.auto_filter.ref and ws.freeze_panes=="A2" and ws.cell(5,2).number_format!="General"); results.append(crit("R005",professional,f"filter={ws.auto_filter.ref if ws else None}; format={ws.cell(5,2).number_format if ws else None}"))
        hash_ok=digest(src)==EXPECTED_SHA256; results.append(crit("R006",hash_ok,f"input_sha256={digest(src)}"))
        if wb: wb.close()
    except Exception as exc: results=[crit("R001",False,f"Judge failure: {exc}")]
    vals={x["criterion_id"]:x["score"] for x in results}; reward=sum(WEIGHTS[k]*vals.get(k,0) for k in WEIGHTS); passed=reward>=.8 and all(vals.get(k)==1.0 for k in ("R001","R002","R003")); payload={"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":[digest(out)] if out.is_file() else [],"criterion_results":results,"reward":reward,"passed":passed}; a.result.parent.mkdir(parents=True,exist_ok=True); a.result.write_text(json.dumps(payload,indent=2),encoding="utf-8"); a.reward_file.parent.mkdir(parents=True,exist_ok=True); a.reward_file.write_text(f"{reward:.6f}\n"); raise SystemExit(0 if passed else 1)
if __name__=="__main__": main()
