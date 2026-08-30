#!/usr/bin/env python3
"""Independent Judge for PDF conversion and MIME attachment semantics."""
from __future__ import annotations
import argparse, hashlib, json, re, shutil, subprocess, tempfile, zipfile
from email import policy
from email.parser import BytesParser
from pathlib import Path
from openpyxl import load_workbook

TASK_ID="cowork-evo/officebench-3-1-task-pdf-email"; RUBRIC_VERSION="officebench-3-1-task-pdf-email-v0.1"; JUDGE_VERSION="officebench-3-1-task-pdf-email-deterministic-v0.1"; EXPECTED_SHA256="4731eca77b9807cd708f298d47d34a75b46c728929e416dc47b0ee530ee8b058"; WEIGHTS={"R001":.15,"R002":.30,"R003":.15,"R004":.25,"R005":.10,"R006":.05}
def digest(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def crit(cid,s,e): return {"criterion_id":cid,"score":float(s),"evidence":e,"method":"deterministic"}
def main():
    ap=argparse.ArgumentParser(); ap.add_argument("--artifact-dir",type=Path,required=True); ap.add_argument("--result",type=Path,required=True); ap.add_argument("--reward-file",type=Path,required=True); a=ap.parse_args(); out=a.artifact_dir; pdf=out/"task.pdf"; eml=out/"Alice.eml"; src=Path(__file__).parents[1]/"data/input_files/task_input.xlsx"; results=[]
    try:
        pdf_ok=pdf.is_file() and pdf.read_bytes().startswith(b"%PDF")
        eml_ok=eml.is_file() and eml.read_bytes().startswith((b"From:",b"MIME-Version:",b"To:"))
        results.append(crit("R001",pdf_ok and eml_ok,f"pdf={pdf_ok}; eml={eml_ok}"))
        source_wb=load_workbook(src,data_only=True); ws=source_wb.active; names=[str(ws.cell(r,4).value) for r in range(3,ws.max_row+1) if ws.cell(r,4).value]; source_wb.close()
        text=""
        if pdf_ok:
            with tempfile.TemporaryDirectory() as tmp:
                txt=Path(tmp)/"task.txt"; subprocess.run(["pdftotext","-layout",str(pdf),str(txt)],check=True,stdout=subprocess.PIPE,stderr=subprocess.PIPE); text=txt.read_text(encoding="utf-8",errors="replace")
        present=[name for name in names if name in text]; all_present=len(present)==len(names); results.append(crit("R002",all_present,f"expected_names={names}; present={present}"))
        occurrence_counts={name: len(re.findall(r"(?<![A-Za-z0-9])" + re.escape(name) + r"(?![A-Za-z0-9])", text)) for name in names}
        occurrence_ok=all(count==1 for count in occurrence_counts.values()) and not any(token in text for token in ["Activaty1","Activaty2","Activaty3","Activaty4"]); results.append(crit("R003",occurrence_ok,f"name_occurrences={occurrence_counts}"))
        msg=BytesParser(policy=policy.default).parsebytes(eml.read_bytes()) if eml_ok else None; attachments=list(msg.iter_attachments()) if msg else []; to_values=[str(x) for x in msg.get_all("To",[]) ] if msg else []; attachment_ok=bool(msg and any(p.get_filename()=="task.pdf" and p.get_content_type()=="application/pdf" and p.get_payload(decode=True)==pdf.read_bytes() for p in attachments)); addressed=any("alice" in v.lower() for v in to_values); results.append(crit("R004",addressed and attachment_ok,f"to={to_values}; attachments={[p.get_filename() for p in attachments]}"))
        body=msg.get_body(preferencelist=("plain",)).get_content() if msg and msg.get_body(preferencelist=("plain",)) else ""; body_ok="review" in body.lower() and "project task" in body.lower(); results.append(crit("R005",body_ok and attachment_ok,f"body={body[:160]!r}; attachment_match={attachment_ok}"))
        hash_ok=digest(src)==EXPECTED_SHA256; results.append(crit("R006",hash_ok,f"input_sha256={digest(src)}"))
    except Exception as exc: results=[crit("R001",False,f"Judge failure: {exc}")]
    vals={x["criterion_id"]:x["score"] for x in results}; reward=sum(WEIGHTS[k]*vals.get(k,0) for k in WEIGHTS); passed=reward>=.8 and all(vals.get(k)==1.0 for k in ("R001","R002","R004")); payload={"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":[digest(p) for p in (pdf,eml) if p.is_file()],"criterion_results":results,"reward":reward,"passed":passed}; a.result.parent.mkdir(parents=True,exist_ok=True); a.result.write_text(json.dumps(payload,indent=2),encoding="utf-8"); a.reward_file.parent.mkdir(parents=True,exist_ok=True); a.reward_file.write_text(f"{reward:.6f}\n"); raise SystemExit(0 if passed else 1)
if __name__=="__main__": main()
