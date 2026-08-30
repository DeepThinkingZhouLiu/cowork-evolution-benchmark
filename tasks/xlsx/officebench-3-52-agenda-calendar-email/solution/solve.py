#!/usr/bin/env python3
"""从会议议程工作簿生成一个日历事件和三封独立提醒邮件。"""

from __future__ import annotations

import argparse
import hashlib
from datetime import datetime, time
from email.message import EmailMessage
from pathlib import Path

from openpyxl import load_workbook

INPUT_NAME = "meeting_agenda.xlsx"
EXPECTED_SHA256 = "ffd3161364fde910fbd5491cc1dd19ba346134231e60f217ad668a874af80b0e"
DATE = "20200501"
DTSTAMP = "20200101T000000Z"


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def seconds(value: object) -> int:
    if isinstance(value, datetime):
        return value.hour * 3600 + value.minute * 60 + value.second
    if isinstance(value, time):
        return value.hour * 3600 + value.minute * 60 + value.second
    if isinstance(value, (int, float)):
        return round(float(value) * 86400)
    raise ValueError(f"无法解析时间：{value!r}")


def hhmmss(value: object) -> str:
    total = seconds(value)
    return f"{total // 3600:02d}{(total % 3600) // 60:02d}{total % 60:02d}"


def escape_ics(value: str) -> str:
    return value.replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")


def make_calendar(rows: list[tuple[object, object, str, str]]) -> str:
    start = hhmmss(rows[0][0])
    end = hhmmss(rows[-1][1])
    agenda = "; ".join(f"{item} ({contact})" for _, _, item, contact in rows)
    lines = [
        "BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//CoworkEvoBench//OfficeBench 3-52//EN",
        "CALSCALE:GREGORIAN", "METHOD:PUBLISH", "BEGIN:VEVENT",
        "UID:bob-agenda-20200501@coworkevobench.local", f"DTSTAMP:{DTSTAMP}",
        f"DTSTART:{DATE}T{start}", f"DTEND:{DATE}T{end}",
        "SUMMARY:Meeting agenda", f"DESCRIPTION:{escape_ics(agenda)}",
        "END:VEVENT", "END:VCALENDAR", "",
    ]
    return "\r\n".join(lines)


def make_email(recipient: str, item: str, contact: str) -> str:
    message = EmailMessage()
    message["From"] = "Bob@example.com"
    message["To"] = f"{recipient}@example.com"
    message["Subject"] = "Meeting agenda reminder"
    message.set_content(f"Meeting agenda reminder for {recipient}. Your agenda item is: {item}. Contact: {contact}.")
    return message.as_string()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    source_path = args.input_dir / INPUT_NAME
    if digest(source_path) != EXPECTED_SHA256:
        raise ValueError("meeting_agenda.xlsx 输入哈希不匹配")
    workbook = load_workbook(source_path, data_only=True, read_only=True)
    if workbook.sheetnames != ["Meeting agenda"]:
        raise ValueError(f"工作表结构异常：{workbook.sheetnames}")
    sheet = workbook["Meeting agenda"]
    headers = [sheet.cell(2, col).value for col in range(2, 7)]
    if headers != ["Start", "End", "Time", "Item", "Contact"]:
        raise ValueError(f"议程表头异常：{headers!r}")
    rows = []
    for values in sheet.iter_rows(min_row=3, max_row=sheet.max_row, min_col=2, max_col=6, values_only=True):
        if values[3] is not None:
            rows.append((values[0], values[1], str(values[3]), str(values[4])))
    workbook.close()
    if rows != [(time(13, 0), time(13, 30), "Introductions", "Linda"), (time(13, 30), time(14, 10), "Agenda", "Davide"), (time(14, 10), time(14, 30), "Strategy", "Charlie")]:
        raise ValueError(f"议程内容与固定输入不符：{rows!r}")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / "Bob.ics").write_text(make_calendar(rows), encoding="utf-8", newline="")
    for start, end, item, contact in rows:
        (args.output_dir / f"{contact}.eml").write_text(make_email(contact, item, contact), encoding="utf-8", newline="")


if __name__ == "__main__":
    main()
