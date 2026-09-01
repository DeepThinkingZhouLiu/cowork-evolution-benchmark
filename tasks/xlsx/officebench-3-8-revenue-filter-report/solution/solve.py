#!/usr/bin/env python3
"""按收入阈值筛选并生成单一 Excel 工作簿。"""
from __future__ import annotations
import argparse
import hashlib
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

INPUT_NAME = "company_revenues.xlsx"
OUTPUT_NAME = "Revenue_Filter.xlsx"
EXPECTED_SHA256 = "b7a5d913221fafb9f545e9367df2bcf07e665b93baefe726788a88532c937e28"
THRESHOLD = 2_000_000

def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()

def style(sheet, widths: dict[str, float], table_name: str) -> None:
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width
    table = Table(displayName=table_name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)

def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    source = args.input_dir / INPUT_NAME
    if digest(source) != EXPECTED_SHA256:
        raise ValueError("company_revenues.xlsx 输入哈希不匹配")
    source_book = load_workbook(source, data_only=False)
    try:
        source_sheet = source_book.active
        rows = [(row[0].value, row[1].value) for row in source_sheet.iter_rows(min_row=2, max_col=2) if row[0].value is not None]
        retained = [row for row in rows if float(row[1]) >= THRESHOLD]
        removed = [row for row in rows if float(row[1]) < THRESHOLD]
        workbook = Workbook()
        filtered = workbook.active
        filtered.title = "Filtered Revenue"
        filtered.append([source_sheet.cell(1, 1).value, source_sheet.cell(1, 2).value])
        for row in retained:
            filtered.append(row)
        for cell in filtered["B"][1:]:
            cell.number_format = "$#,##0"
        style(filtered, {"A": 14, "B": 20}, "Filtered_Revenue")
        summary = workbook.create_sheet("Filter Summary")
        summary.append(["Metric", "Value"])
        for row in [("Threshold", THRESHOLD), ("Input Rows", len(rows)), ("Retained Rows", len(retained)), ("Deleted Rows", len(removed))]:
            summary.append(row)
        summary["B2"].number_format = "$#,##0"
        style(summary, {"A": 20, "B": 18}, "Filter_Summary")
        deleted = workbook.create_sheet("Deleted Records")
        deleted.append(["Year", "Revenue ($)"])
        for row in removed:
            deleted.append(row)
        for cell in deleted["B"][1:]:
            cell.number_format = "$#,##0"
        style(deleted, {"A": 14, "B": 20}, "Deleted_Records")
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
        args.output_dir.mkdir(parents=True, exist_ok=True)
        workbook.save(args.output_dir / OUTPUT_NAME)
        workbook.close()
    finally:
        source_book.close()

if __name__ == "__main__":
    main()
