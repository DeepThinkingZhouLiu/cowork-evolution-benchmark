#!/usr/bin/env python3
"""从输入独立重算收入筛选和摘要工作表。"""
from __future__ import annotations
import argparse
import hashlib
import json
import zipfile
from pathlib import Path
from openpyxl import load_workbook

TASK_ID = "cowork-evo/officebench-3-8-revenue-filter-report"
RUBRIC_VERSION = "officebench-3-8-revenue-filter-report-v0.2"
JUDGE_VERSION = "officebench-3-8-revenue-filter-report-deterministic-v0.2"
INPUT_NAME = "company_revenues.xlsx"
OUTPUT_NAME = "Revenue_Filter.xlsx"
THRESHOLD = 2_000_000
EXPECTED_SHA256 = "b7a5d913221fafb9f545e9367df2bcf07e665b93baefe726788a88532c937e28"
WEIGHTS = {"R001": .15, "R002": .35, "R003": .10, "R004": .15, "R005": .15, "R006": .05, "R007": .05}

def criterion(cid: str, ok: bool, evidence: str) -> dict:
    return {"criterion_id": cid, "score": float(ok), "evidence": evidence, "method": "deterministic"}

def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()

def records(sheet) -> list[tuple[int, int]]:
    return [(int(row[0]), int(row[1])) for row in sheet.iter_rows(min_row=2, max_col=2, values_only=True) if row[0] is not None]

def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    root = Path(__file__).parents[1]
    source_path = root / "data/input_files" / INPUT_NAME
    output_path = args.artifact_dir / OUTPUT_NAME
    results: list[dict] = []
    try:
        source = load_workbook(source_path, data_only=True)
        source_rows = records(source.active)
        source.close()
        expected = [row for row in source_rows if row[1] >= THRESHOLD]
        removed = [row for row in source_rows if row[1] < THRESHOLD]
        valid = output_path.is_file() and zipfile.is_zipfile(output_path)
        if not valid:
            results = [criterion("R001", False, "Revenue_Filter.xlsx 缺失或不是有效 OOXML")]
            results.extend(criterion(cid, False, "缺少有效交付") for cid in ("R002", "R003", "R004", "R005", "R006"))
            results.append(criterion("R007", digest(source_path) == EXPECTED_SHA256, "输入哈希检查"))
        else:
            workbook = load_workbook(output_path, data_only=False)
            structure = workbook.sheetnames == ["Filtered Revenue", "Filter Summary", "Deleted Records"]
            results.append(criterion("R001", structure, f"工作表={workbook.sheetnames}"))
            filtered = workbook["Filtered Revenue"] if "Filtered Revenue" in workbook.sheetnames else workbook.active
            filtered_rows = records(filtered)
            headers_ok = tuple(cell.value for cell in filtered[1][:2]) == ("Year", "Revenue ($)")
            values_ok = filtered_rows == expected and headers_ok and filtered.max_column == 2
            results.append(criterion("R002", values_ok, f"保留行={len(filtered_rows)}，期望={len(expected)}，顺序和值一致={values_ok}"))
            invalid = [row for row in filtered_rows if row[1] < THRESHOLD]
            unique_ok = len(filtered_rows) == len(set(filtered_rows)) and not invalid and set(filtered_rows) == set(expected)
            results.append(criterion("R003", unique_ok, f"低于阈值={invalid}，无重复/额外={unique_ok}"))
            summary = workbook["Filter Summary"] if "Filter Summary" in workbook.sheetnames else None
            summary_values = {str(row[0]).strip(): row[1] for row in summary.iter_rows(min_row=2, max_col=2, values_only=True) if row[0] is not None} if summary else {}
            expected_summary = {"Threshold": THRESHOLD, "Input Rows": len(source_rows), "Retained Rows": len(expected), "Deleted Rows": len(removed)}
            results.append(criterion("R004", summary_values == expected_summary, f"筛选摘要={summary_values}"))
            deleted = workbook["Deleted Records"] if "Deleted Records" in workbook.sheetnames else None
            deleted_rows = records(deleted) if deleted else []
            deleted_ok = bool(deleted) and deleted_rows == removed and tuple(cell.value for cell in deleted[1][:2]) == ("Year", "Revenue ($)")
            results.append(criterion("R005", deleted_ok, f"删除记录={deleted_rows}"))
            professional = all(workbook[name].freeze_panes == "A2" and workbook[name].auto_filter.ref for name in workbook.sheetnames)
            editable = all(workbook[name].protection.sheet is not True for name in workbook.sheetnames)
            results.append(criterion("R006", professional and editable, f"筛选冻结={professional}，可编辑={editable}"))
            workbook.close()
            results.append(criterion("R007", digest(source_path) == EXPECTED_SHA256, f"输入 SHA256={digest(source_path)}"))
    except Exception as exc:
        results = [criterion("R001", False, f"Judge 读取失败：{exc}")] + [criterion(cid, False, "Judge 读取失败") for cid in ("R002", "R003", "R004", "R005", "R006", "R007")]
    values = {item["criterion_id"]: item["score"] for item in results}
    reward = sum(WEIGHTS[cid] * values.get(cid, 0.0) for cid in WEIGHTS)
    passed = all(values.get(cid, 0.0) == 1.0 for cid in ("R001", "R002", "R003", "R004", "R005")) and reward >= .7
    payload = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [digest(output_path)] if output_path.is_file() else [], "criterion_results": results, "reward": reward, "passed": passed}
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.write_text(f"{reward:.6f}\n", encoding="utf-8")
    raise SystemExit(0 if passed else 1)

if __name__ == "__main__":
    main()
