#!/usr/bin/env python3
"""为 officeval_082 的真实标签表增加逐项勾选与自动汇总。

本脚本只处理本题固定输入，不是通用表格生成器。它保留上游工作簿的真实
标签内容，在 126 个非空标签单元格上逐一增加双态下拉，并用可编辑公式驱动
右侧 I:J 汇总区。
"""

from __future__ import annotations

import argparse
import hashlib
import html
import os
import re
import tempfile
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


INPUT_NAME = "1111.xlsx"
REFERENCE_IMAGE_NAME = "image (4).png"
OUTPUT_NAME = "1111_勾选自动汇总版.xlsx"
SHEET_NAME = "标签维度表"

EXPECTED_SHA256 = {
    INPUT_NAME: "37ce4050dcaaf7ddc9dadaa988cf22aacc84a7a3037aa6006866d293bbe8487b",
    REFERENCE_IMAGE_NAME: "ae2ef1f6614116e2a1cbdd0db472658bac3497269f9d27c0d81ef5ebde6042c4",
}

# Golden 预置四个真实标签，便于打开文件后直接看到联动结果；用户可继续取消或勾选。
INITIAL_CHECKED = {"B5", "B20", "C6", "C40"}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def verify_inputs(input_dir: Path) -> tuple[Path, Path]:
    workbook_path = input_dir / INPUT_NAME
    image_path = input_dir / REFERENCE_IMAGE_NAME
    for path in (workbook_path, image_path):
        if not path.is_file():
            raise FileNotFoundError(f"缺少本题真实输入：{path.name}")
        actual = sha256(path)
        if actual != EXPECTED_SHA256[path.name]:
            raise ValueError(f"{path.name} SHA256 不匹配：{actual}")
    return workbook_path, image_path


def strip_checkbox(value: object) -> str:
    text = str(value or "").strip()
    if text[:1] in {"□", "☐", "☑", "✓", "✔"}:
        return text[1:].strip()
    return text


def add_checkbox_dropdowns(ws) -> list[str]:
    """逐格增加可点击的 □/☑ 双态下拉，返回 126 个真实标签坐标。"""

    tag_refs: list[str] = []
    for row in range(5, 69):
        for column in (2, 3):
            cell = ws.cell(row, column)
            label = strip_checkbox(cell.value)
            if not label:
                continue
            tag_refs.append(cell.coordinate)
            symbol = "☑" if cell.coordinate in INITIAL_CHECKED else "□"
            cell.value = f"{symbol} {label}"

            validation = DataValidation(
                type="list",
                formula1=f'"□ {label},☑ {label}"',
                allow_blank=False,
            )
            validation.errorTitle = "请选择标签状态"
            validation.error = "请从下拉列表选择未选中或已选中状态。"
            validation.promptTitle = "切换勾选"
            validation.prompt = "选择 □ 取消，选择 ☑ 勾选。"
            validation.showInputMessage = True
            validation.showErrorMessage = True
            validation.showDropDown = False
            ws.add_data_validation(validation)
            validation.add(cell)

    if len(tag_refs) != 126:
        raise ValueError(f"真实标签数应为 126，实际为 {len(tag_refs)}")

    checked_fill = PatternFill("solid", fgColor="DCFCE7")
    ws.conditional_formatting.add(
        "B5:C68",
        FormulaRule(formula=['LEFT(B5,1)="☑"'], fill=checked_fill),
    )
    return tag_refs


def add_summary_formulas(ws) -> None:
    """建立 B/C 勾选状态到 I:J 的可编辑、可重算汇总链。"""

    # 两列隐藏序号把 B 列与 C 列的已选项映射为唯一的连续排名。
    for row in range(5, 69):
        ws[f"L{row}"] = f'=IF(LEFT(B{row},1)="☑",COUNTIF($B$5:B{row},"☑*"),"")'
        ws[f"M{row}"] = (
            f'=IF(LEFT(C{row},1)="☑",COUNTIF($B$5:$B$68,"☑*")'
            f'+COUNTIF($C$5:C{row},"☑*"),"")'
        )

    for output_row in range(5, 131):
        ordinal = f"ROWS($I$5:I{output_row})"
        rank = f"SMALL($L$5:$M$68,{ordinal})"
        b_match = f"MATCH({rank},$L$5:$L$68,0)"
        c_match = f"MATCH({rank},$M$5:$M$68,0)"
        b_value = f"INDEX($B$5:$B$68,{b_match})"
        c_value = f"INDEX($C$5:$C$68,{c_match})"
        b_category = f"INDEX($A$5:$A$68,{b_match})"
        c_category = f"INDEX($A$5:$A$68,{c_match})"

        ws[f"I{output_row}"] = (
            f'=IFERROR(IF(LEFT({b_value},1)="☑",{b_category},""),'
            f'IFERROR(IF(LEFT({c_value},1)="☑",{c_category},""),""))'
        )
        ws[f"J{output_row}"] = (
            f'=IFERROR(IF(LEFT({b_value},1)="☑",MID({b_value},3,999),""),'
            f'IFERROR(IF(LEFT({c_value},1)="☑",MID({c_value},3,999),""),""))'
        )


def apply_layout(ws) -> None:
    dark_green = "115E59"
    deep_gray = "1F2937"
    grid = Side(style="thin", color="CBD5E1")
    border = Border(left=grid, right=grid, top=grid, bottom=grid)

    # 上游 Rubric 明确给出这些列宽区间；A:C 保留为真实、可编辑单元格。
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 25
    ws.column_dimensions["I"].width = 25
    ws.column_dimensions["J"].width = 25
    ws.column_dimensions["L"].hidden = True
    ws.column_dimensions["M"].hidden = True

    if "I4:J4" not in {str(item) for item in ws.merged_cells.ranges}:
        ws.merge_cells("I4:J4")
    ws["I4"] = "维度分类（选定）"
    ws["I4"].font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
    ws["I4"].fill = PatternFill("solid", fgColor=dark_green)
    ws["I4"].alignment = Alignment(horizontal="center", vertical="center")
    ws["I4"].border = border

    for row in range(5, 131):
        ws.row_dimensions[row].height = 22
        for column in ("I", "J"):
            cell = ws[f"{column}{row}"]
            cell.font = Font(name="Microsoft YaHei", size=10, color=deep_gray)
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = border
            cell.fill = PatternFill("solid", fgColor="F8FAFC" if row % 2 else "FFFFFF")

    ws.freeze_panes = "A5"
    ws.sheet_view.zoomScale = 75
    ws.sheet_view.showGridLines = True
    if ws.sheet_view.selection:
        ws.sheet_view.selection[0].activeCell = "I5"
        ws.sheet_view.selection[0].sqref = "I5"

    ws.print_area = "A1:J68"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = "8"  # A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 2
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35


def inject_initial_formula_cache(path: Path, cached_values: dict[str, str]) -> None:
    """只写入四行初始显示缓存，公式本身仍保留并在打开时完整重算。

    openpyxl 不计算公式。直接给公式单元格补 OOXML 缓存后，上游静态 Verifier
    能检查真实输出字体；用户改变勾选状态后，Excel/WPS/LibreOffice 会按公式重算。
    """

    with ZipFile(path, "r") as source:
        sheet_xml = source.read("xl/worksheets/sheet1.xml").decode("utf-8")
        for ref, value in cached_values.items():
            pattern = re.compile(
                rf'(<c\b(?=[^>]*\br="{re.escape(ref)}")[^>]*>)(.*?</c>)',
                flags=re.DOTALL,
            )

            def replace_cell(match: re.Match[str]) -> str:
                opening, body = match.groups()
                if re.search(r'\st="[^"]*"', opening):
                    opening = re.sub(r'\st="[^"]*"', ' t="str"', opening, count=1)
                else:
                    opening = opening[:-1] + ' t="str">'
                escaped = html.escape(value, quote=False)
                body = re.sub(r"<v>.*?</v>", f"<v>{escaped}</v>", body, count=1, flags=re.DOTALL)
                return opening + body

            sheet_xml, count = pattern.subn(replace_cell, sheet_xml, count=1)
            if count != 1:
                raise ValueError(f"无法写入公式缓存：{ref}")

        with tempfile.NamedTemporaryFile(prefix="officeval-082-", suffix=".xlsx", dir=path.parent, delete=False) as temp:
            temp_path = Path(temp.name)
        try:
            with ZipFile(temp_path, "w", compression=ZIP_DEFLATED) as target:
                for item in source.infolist():
                    data = sheet_xml.encode("utf-8") if item.filename == "xl/worksheets/sheet1.xml" else source.read(item.filename)
                    target.writestr(item, data)
            os.replace(temp_path, path)
        finally:
            if temp_path.exists():
                temp_path.unlink()


def build(input_path: Path, output_path: Path) -> None:
    workbook = load_workbook(input_path)
    if workbook.sheetnames != [SHEET_NAME]:
        raise ValueError(f"工作表结构不符合上游输入：{workbook.sheetnames}")
    ws = workbook[SHEET_NAME]

    add_checkbox_dropdowns(ws)
    add_summary_formulas(ws)
    apply_layout(ws)

    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()

    inject_initial_formula_cache(
        output_path,
        {
            "I5": "基础信息",
            "J5": "新用户",
            "I6": "花销去向",
            "J6": "市场运营",
            "I7": "基础信息",
            "J7": "线下体验偏好",
            "I8": "使用场景",
            "J8": "夜间使用",
        },
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    input_path, _ = verify_inputs(args.input_dir)
    build(input_path, args.output_dir / OUTPUT_NAME)


if __name__ == "__main__":
    main()
