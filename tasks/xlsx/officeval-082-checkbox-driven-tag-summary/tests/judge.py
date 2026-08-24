#!/usr/bin/env python3
"""officeval_082 的 Harbor Judge。

除结构与样式检查外，本 Judge 会把两组不同标签实际切换为 ☑，再调用
LibreOffice 重算，确认右侧只出现对应分类与标签。这样可以区分真正可 replay
的联动工作簿和静态写死的“标准答案”。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import subprocess
import tempfile
import warnings
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import upstream_verifier


OUTPUT_NAME = "1111_勾选自动汇总版.xlsx"
INPUT_NAME = "1111.xlsx"
SHEET_NAME = "标签维度表"
TASK_ID = "cowork-evo/officeval-082-checkbox-driven-tag-summary"
RUBRIC_VERSION = "omegause-officeval-082-adaptation-v0.1"
JUDGE_VERSION = "officeval-082-behavioral-v0.1"
BAD_OUTPUTS = {"TRUE", "FALSE", "0", "1", "#VALUE!", "#N/A", "#CALC!", "#REF!", "#DIV/0!", "#NAME?"}
CHECK_SYMBOLS = {"□", "☐", "☑", "✓", "✔"}

CRITERION_SPECS = {
    "R001": ("指定 XLSX 可打开且只含原标签维度表", 2, "hurdle", "deterministic"),
    "R002": ("五类 126 个真实标签及左侧可编辑结构完整保留", 5, "positive", "deterministic"),
    "R003": ("126 个标签逐项具有可切换且一一对应的勾选控件", 6, "positive", "deterministic"),
    "R004": ("B/C 选中状态与 I:J 动态汇总公式完整关联", 6, "positive", "deterministic"),
    "R005": ("右侧标题、列宽、行高、字体、边框及显示值符合要求", 4, "positive", "deterministic"),
    "R006": ("公式、控件和工作表均可编辑并启用自动重算", 3, "positive", "deterministic"),
    "R007": ("两组独立勾选场景经 LibreOffice 重算均得到精确结果", 8, "positive", "behavioral_replay"),
    "R008": ("LibreOffice 可正常渲染完整左表和右侧汇总区", 3, "positive", "deterministic"),
    "R009": ("真实标签缺失、错改或勾选控件不完整", -10, "penalty", "deterministic"),
    "R010": ("右侧结果静态写死、只关联单列或实际回放失败", -10, "penalty", "behavioral_replay"),
    "R011": ("左表被图片替代、工作表受保护或无法继续编辑", -8, "penalty", "deterministic"),
    "R012": ("关键列宽、标题、显示值或实际渲染不符合", -6, "penalty", "deterministic"),
}


def criterion(identifier: str, passed: bool, evidence: str = "") -> dict:
    description, weight, kind, method = CRITERION_SPECS[identifier]
    return {
        "criterion_id": identifier,
        "description": description,
        "weight": weight,
        "type": kind,
        "passed": bool(passed),
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": method,
    }


def strip_checkbox(value: object) -> str:
    text = str(value or "").strip()
    if text[:1] in CHECK_SYMBOLS:
        return text[1:].strip()
    return text


def source_tag_refs(source) -> list[str]:
    refs = []
    for row in range(5, 69):
        for column in ("B", "C"):
            if str(source[f"{column}{row}"].value or "").strip():
                refs.append(f"{column}{row}")
    return refs


def source_fidelity_errors(source, candidate) -> list[str]:
    errors: list[str] = []
    refs = set(source_tag_refs(source))
    if len(refs) != 126:
        errors.append(f"基准输入标签数异常：{len(refs)}")

    for row in range(1, 69):
        for column in range(1, 4):
            before = source.cell(row, column)
            after = candidate.cell(row, column)
            if before.coordinate in refs:
                if strip_checkbox(after.value) != strip_checkbox(before.value):
                    errors.append(f"{before.coordinate} 标签：{before.value!r} -> {after.value!r}")
                current = str(after.value or "").strip()
                if current[:1] not in CHECK_SYMBOLS:
                    errors.append(f"{before.coordinate} 缺少可见勾选状态前缀")
            elif before.value != after.value and not (
                before.value in (None, "") and after.value in (None, "")
            ):
                errors.append(f"{before.coordinate}：{before.value!r} -> {after.value!r}")
            # openpyxl 保存时会重建合并单元格的从属格样式，也会归一化第 3 行的
            # 空白占位样式；这些格没有可编辑内容。真实左表标题、表头和 64 行主体
            # 的样式必须逐格保持。
            style_must_match = before.coordinate == "A1" or row >= 4
            if style_must_match and before.style_id != after.style_id:
                errors.append(f"{before.coordinate} 原样式改变：{before.style_id} -> {after.style_id}")

    source_merges = {str(item) for item in source.merged_cells.ranges}
    candidate_merges = {str(item) for item in candidate.merged_cells.ranges}
    if not source_merges.issubset(candidate_merges):
        errors.append(f"原合并区域缺失：{sorted(source_merges - candidate_merges)}")
    return errors


def validation_coverage(ws, source) -> tuple[list[str], int]:
    expected_refs = source_tag_refs(source)
    by_cell: dict[str, list[str]] = {}
    for validation in ws.data_validations.dataValidation:
        if validation.type != "list":
            continue
        formula = str(validation.formula1 or "")
        for cell_range in validation.ranges.ranges:
            for row, column in cell_range.cells:
                ref = f"{get_column_letter(column)}{row}"
                by_cell.setdefault(ref, []).append(formula)

    errors = []
    for ref in expected_refs:
        label = strip_checkbox(source[ref].value)
        formulas = by_cell.get(ref, [])
        matching = [item for item in formulas if "□" in item and "☑" in item and label in item]
        if len(matching) != 1:
            errors.append(f"{ref} 双态控件={len(matching)}")

    unrelated = sorted(set(by_cell) - set(expected_refs))
    if unrelated:
        errors.append(f"非标签单元格存在勾选验证：{unrelated[:10]}")
    return errors, len(by_cell)


def formula_and_editability_errors(workbook, ws) -> list[str]:
    errors = []
    formulas = []
    output_formulas = []
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formulas.append(cell.value)
                if cell.column in (9, 10) and 5 <= cell.row <= 200:
                    output_formulas.append(cell.value)
    joined = " ".join(output_formulas).upper()
    if len(output_formulas) < 2:
        errors.append(f"I:J 动态公式仅 {len(output_formulas)} 个")
    if not re.search(r"\$?B\$?\d+", joined):
        errors.append("右侧公式未引用 B 列")
    if not re.search(r"\$?C\$?\d+", joined):
        errors.append("右侧公式未引用 C 列")
    if "☑" not in joined or "LEFT(" not in joined:
        errors.append("右侧公式未按 ☑ 状态过滤")
    if not any(token in joined for token in ("SMALL(", "AGGREGATE(", "FILTER(")):
        errors.append("右侧公式没有可保留多选项的枚举逻辑")
    if ws.protection.sheet is True:
        errors.append("工作表已保护")
    calc = workbook.calculation
    if calc.calcMode != "auto" or not bool(calc.fullCalcOnLoad) or not bool(calc.forceFullCalc):
        errors.append(
            f"重算设置={calc.calcMode}/{calc.fullCalcOnLoad}/{calc.forceFullCalc}"
        )
    return errors


def presentation_errors(ws, data_only_ws) -> list[str]:
    errors = []
    merges = {str(item) for item in ws.merged_cells.ranges}
    if ws["I4"].value != "维度分类（选定）" or "I4:J4" not in merges:
        errors.append(f"右侧标题={ws['I4'].value!r}，合并={sorted(merges)}")

    expected_widths = {"A": (8, 12), "B": (24, 26), "C": (24, 26), "I": (24, 26), "J": (24, 26)}
    for column, bounds in expected_widths.items():
        width = ws.column_dimensions[column].width
        if width is None or not bounds[0] <= float(width) <= bounds[1]:
            errors.append(f"{column} 列宽={width}，期望 {bounds[0]}–{bounds[1]}")

    output_formula_rows = sorted(
        {
            cell.row
            for row in ws.iter_rows(min_col=9, max_col=10, min_row=5, max_row=200)
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        }
    )
    if not output_formula_rows:
        errors.append("右侧没有公式行")
    for row in output_formula_rows:
        height = ws.row_dimensions[row].height
        if height is None or not 20 <= float(height) <= 24:
            errors.append(f"第 {row} 行行高={height}")
            break

    accepted_fonts = {"微软雅黑", "Microsoft YaHei", "Microsoft YaHei UI", "宋体", "SimSun", "Calibri", "Carlito"}
    for row in output_formula_rows[:20]:
        for column in (9, 10):
            cell = ws.cell(row, column)
            color = cell.font.color
            rgb = color.rgb[-6:].upper() if color is not None and color.type == "rgb" and color.rgb else "000000"
            if cell.font.name not in accepted_fonts or not (9 <= float(cell.font.sz or 0) <= 12):
                errors.append(f"{cell.coordinate} 字体={cell.font.name}/{cell.font.sz}")
            if rgb not in {"000000", "111827", "1F2937", "333333", "404040", "0F172A"}:
                errors.append(f"{cell.coordinate} 字色=#{rgb}")
            if cell.alignment.horizontal not in (None, "general", "left", "center"):
                errors.append(f"{cell.coordinate} 对齐={cell.alignment.horizontal}")
        if len(errors) > 20:
            break

    for row in range(5, 201):
        for column in ("I", "J"):
            value = data_only_ws[f"{column}{row}"].value
            if value is not None and str(value).strip().upper() in BAD_OUTPUTS:
                errors.append(f"{column}{row} 显示中间值 {value!r}")
    return errors


def contains_embedded_image(path: Path) -> tuple[bool, list[str]]:
    try:
        with ZipFile(path) as archive:
            media = [name for name in archive.namelist() if name.startswith("xl/media/")]
    except Exception:
        return True, ["无法检查 OOXML 媒体"]
    return bool(media), media


def selected_pairs(source, refs: list[str]) -> list[tuple[str, str]]:
    pairs = []
    for ref in refs:
        row = source[ref].row
        pairs.append((str(source[f"A{row}"].value or "").strip(), strip_checkbox(source[ref].value)))
    return pairs


def read_output_pairs(ws) -> tuple[list[tuple[str, str]], list[str]]:
    pairs = []
    errors = []
    for row in range(5, 201):
        category = ws[f"I{row}"].value
        label = ws[f"J{row}"].value
        category_text = str(category or "").strip()
        label_text = str(label or "").strip()
        if not category_text and not label_text:
            continue
        if not category_text or not label_text:
            errors.append(f"I{row}:J{row} 输出不成对：{category!r}/{label!r}")
            continue
        if category_text.upper() in BAD_OUTPUTS or label_text.upper() in BAD_OUTPUTS:
            errors.append(f"I{row}:J{row} 出现中间值：{category!r}/{label!r}")
            continue
        pairs.append((category_text, label_text))
    return pairs, errors


def replay_one(candidate_path: Path, source, selected: list[str], scenario_index: int) -> tuple[bool, str]:
    libreoffice = shutil.which("libreoffice")
    if libreoffice is None:
        return False, "缺少 libreoffice"

    with tempfile.TemporaryDirectory(prefix=f"cowork-evo-082-replay-{scenario_index}-") as temp_dir:
        temp_root = Path(temp_dir)
        mutated_path = temp_root / "scenario.xlsx"
        calc_dir = temp_root / "calculated"
        profile_dir = temp_root / "lo-profile"
        calc_dir.mkdir()
        profile_dir.mkdir()

        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            workbook = load_workbook(candidate_path, data_only=False)
        ws = workbook[SHEET_NAME]
        selected_set = set(selected)
        for ref in source_tag_refs(source):
            label = strip_checkbox(source[ref].value)
            ws[ref] = f"{'☑' if ref in selected_set else '□'} {label}"
        workbook.calculation.calcMode = "auto"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.save(mutated_path)
        workbook.close()

        converted = subprocess.run(
            [
                libreoffice,
                "--headless",
                f"-env:UserInstallation={profile_dir.resolve().as_uri()}",
                "--convert-to",
                "xlsx",
                "--outdir",
                str(calc_dir),
                str(mutated_path),
            ],
            capture_output=True,
            text=True,
            timeout=300,
            check=False,
        )
        calculated_path = calc_dir / mutated_path.name
        if converted.returncode != 0 or not calculated_path.is_file():
            return False, (converted.stderr or converted.stdout or "LibreOffice 重算失败")[-1000:]

        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            calculated = load_workbook(calculated_path, data_only=True, read_only=True)
        calculated_ws = calculated[SHEET_NAME]
        actual, errors = read_output_pairs(calculated_ws)
        calculated.close()

    expected = selected_pairs(source, selected)
    ok = not errors and Counter(actual) == Counter(expected) and len(actual) == len(expected)
    return ok, f"勾选={selected}，期望={expected}，实际={actual}，错误={errors}"


def replay_scenarios(candidate_path: Path, source, scenarios: list[list[str]]) -> tuple[bool, list[str]]:
    evidence = []
    all_ok = True
    for index, selected in enumerate(scenarios, start=1):
        ok, detail = replay_one(candidate_path, source, selected, index)
        all_ok = all_ok and ok
        evidence.append(f"场景{index}：通过={ok}；{detail}")
    return all_ok, evidence


def render_check(path: Path) -> tuple[bool, str]:
    libreoffice = shutil.which("libreoffice")
    pdfinfo = shutil.which("pdfinfo")
    pdftotext = shutil.which("pdftotext")
    if libreoffice is None or pdfinfo is None or pdftotext is None:
        return False, "缺少 libreoffice、pdfinfo 或 pdftotext"
    with tempfile.TemporaryDirectory(prefix="cowork-evo-082-render-") as temp_dir:
        temp_root = Path(temp_dir)
        profile_dir = temp_root / "lo-profile"
        profile_dir.mkdir()
        converted = subprocess.run(
            [
                libreoffice,
                "--headless",
                f"-env:UserInstallation={profile_dir.resolve().as_uri()}",
                "--convert-to",
                "pdf",
                "--outdir",
                temp_dir,
                str(path),
            ],
            capture_output=True,
            text=True,
            timeout=300,
            check=False,
        )
        pdf_path = temp_root / f"{path.stem}.pdf"
        if converted.returncode != 0 or not pdf_path.is_file():
            return False, (converted.stderr or converted.stdout or "PDF 转换失败")[-1000:]
        info = subprocess.run([pdfinfo, str(pdf_path)], capture_output=True, text=True, timeout=120, check=False)
        text = subprocess.run([pdftotext, "-layout", str(pdf_path), "-"], capture_output=True, text=True, timeout=120, check=False)
        if info.returncode != 0 or text.returncode != 0:
            return False, (info.stderr or text.stderr or "PDF 读取失败")[-1000:]
        match = re.search(r"^Pages:\s+(\d+)", info.stdout, flags=re.MULTILINE)
        pages = int(match.group(1)) if match else 0
        rendered = re.sub(r"\s+", "", text.stdout)
    required = ["用户标签维度表", "分类维度", "标签选项一", "标签选项二", "维度分类（选定）", "购买动因", "可扩展升级"]
    missing = [item for item in required if re.sub(r"\s+", "", item) not in rendered]
    return 1 <= pages <= 8 and not missing, f"页数={pages}，渲染缺失={missing}，提取字符={len(rendered)}"


def evaluate(artifact_dir: Path) -> dict:
    output_path = artifact_dir / OUTPUT_NAME
    results: list[dict] = []
    if not output_path.is_file():
        results.append(criterion("R001", False, f"缺少 {OUTPUT_NAME}"))
        return build_result(results, output_path)

    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            workbook = load_workbook(output_path, data_only=False)
            cached_workbook = load_workbook(output_path, data_only=True)
    except Exception as exc:
        results.append(criterion("R001", False, f"文件无法打开：{exc}"))
        return build_result(results, output_path)

    sheet_ok = workbook.sheetnames == [SHEET_NAME]
    results.append(criterion("R001", sheet_ok, f"工作表={workbook.sheetnames}"))
    if not sheet_ok:
        workbook.close()
        cached_workbook.close()
        return build_result(results, output_path)

    task_root = Path(__file__).resolve().parents[1]
    expectations = json.loads(Path(__file__).with_name("expectations.json").read_text(encoding="utf-8"))
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        source_workbook = load_workbook(task_root / "data" / "input_files" / INPUT_NAME, data_only=False)
    source = source_workbook[SHEET_NAME]
    ws = workbook[SHEET_NAME]
    cached_ws = cached_workbook[SHEET_NAME]

    fidelity_issues = source_fidelity_errors(source, ws)
    results.append(
        criterion("R002", not fidelity_issues, "5 类 126 个标签及原样式完整" if not fidelity_issues else "；".join(fidelity_issues[:30]))
    )

    validation_issues, validation_cells = validation_coverage(ws, source)
    try:
        upstream = upstream_verifier.evaluate(str(artifact_dir))
    except Exception as exc:
        upstream = {"status": "error", "error": repr(exc), "dim2_items": [], "total_score": 0, "max_score": 0}
    upstream_items = upstream.get("dim2_items") or []
    upstream_controls = bool(len(upstream_items) >= 1 and upstream_items[0].get("hit"))
    controls_ok = not validation_issues and validation_cells == expectations["tag_count"] and upstream_controls
    results.append(
        criterion(
            "R003",
            controls_ok,
            f"逐格双态验证={validation_cells}/126，上游控件判据={upstream_controls}"
            if controls_ok
            else f"验证问题={validation_issues[:20]}，逐格覆盖={validation_cells}，上游={upstream_controls}",
        )
    )

    formula_issues = formula_and_editability_errors(workbook, ws)
    upstream_interaction = bool(len(upstream_items) >= 2 and upstream_items[1].get("hit"))
    formula_ok = not formula_issues and upstream_interaction
    results.append(
        criterion(
            "R004",
            formula_ok,
            "B/C 勾选与 I:J 多选枚举逻辑完整" if formula_ok else f"公式问题={formula_issues}，上游联动={upstream_interaction}",
        )
    )

    layout_issues = presentation_errors(ws, cached_ws)
    upstream_visual_ok = bool(
        len(upstream_items) == 6 and all(item.get("hit") for item in upstream_items[2:])
    )
    layout_ok = not layout_issues and upstream_visual_ok
    results.append(
        criterion(
            "R005",
            layout_ok,
            f"列宽、行高、字体、边框和显示值通过；上游={upstream.get('total_score')}/{upstream.get('max_score')}"
            if layout_ok
            else f"版式问题={layout_issues[:20]}，上游视觉={upstream_visual_ok}",
        )
    )

    embedded_image, media = contains_embedded_image(output_path)
    editable_ok = not formula_issues and ws.protection.sheet is not True and not embedded_image
    results.append(
        criterion(
            "R006",
            editable_ok,
            "公式、逐格下拉和自动重算均可继续编辑" if editable_ok else f"公式问题={formula_issues}，保护={ws.protection.sheet}，媒体={media}",
        )
    )

    replay_ok, replay_evidence = replay_scenarios(output_path, source, expectations["replay_scenarios"])
    results.append(criterion("R007", replay_ok, "；".join(replay_evidence)))

    rendered_ok, render_evidence = render_check(output_path)
    results.append(criterion("R008", rendered_ok, render_evidence))

    results.append(criterion("R009", bool(fidelity_issues or validation_issues or not upstream_controls), f"标签问题={len(fidelity_issues)}，控件问题={len(validation_issues)}"))
    results.append(criterion("R010", bool(not formula_ok or not replay_ok), f"公式通过={formula_ok}，行为回放={replay_ok}"))
    results.append(criterion("R011", bool(embedded_image or ws.protection.sheet is True), f"嵌入媒体={media}，保护={ws.protection.sheet}"))
    results.append(criterion("R012", bool(layout_issues or not upstream_visual_ok or not rendered_ok), f"版式问题={len(layout_issues)}，上游视觉={upstream_visual_ok}，渲染={rendered_ok}"))

    source_workbook.close()
    workbook.close()
    cached_workbook.close()
    return build_result(results, output_path)


def build_result(results: list[dict], artifact_path: Path) -> dict:
    present = {item["criterion_id"] for item in results}
    for identifier in CRITERION_SPECS:
        if identifier not in present:
            results.append(criterion(identifier, False, "因交付不存在、损坏或门槛失败，本判据未执行"))
    results.sort(key=lambda item: item["criterion_id"])
    positive_total = sum(item["weight"] for item in results if item["weight"] > 0)
    earned = sum(item["weight"] for item in results if item["weight"] > 0 and item["passed"])
    penalties = sum(item["weight"] for item in results if item["weight"] < 0 and item["passed"])
    reward = max(0.0, min(1.0, (earned + penalties) / positive_total))
    hurdle_failures = [item["criterion_id"] for item in results if item["type"] == "hurdle" and not item["passed"]]
    if hurdle_failures:
        reward = 0.0
    digest = hashlib.sha256(artifact_path.read_bytes()).hexdigest() if artifact_path.is_file() else hashlib.sha256(b"").hexdigest()
    normalized_results = []
    for item in results:
        score = (0.0 if item["passed"] else 1.0) if item["type"] == "penalty" else item["score"]
        normalized_results.append(
            {
                "criterion_id": item["criterion_id"],
                "score": score,
                "evidence": item["evidence"],
                "method": item["method"],
                "raw": {
                    "description": item["description"],
                    "weight": item["weight"],
                    "type": item["type"],
                    "passed": item["passed"],
                },
            }
        )
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [digest],
        "criterion_results": normalized_results,
        "reward": round(reward, 6),
        "passed": reward >= 0.75,
        "hurdle_failures": hurdle_failures,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path)
    args = parser.parse_args()
    result = evaluate(args.artifact_dir)
    rendered = json.dumps(result, ensure_ascii=False, indent=2)
    print(rendered)
    if args.result:
        args.result.parent.mkdir(parents=True, exist_ok=True)
        args.result.write_text(rendered + "\n", encoding="utf-8")
    raise SystemExit(0 if result["passed"] else 1)


if __name__ == "__main__":
    main()
