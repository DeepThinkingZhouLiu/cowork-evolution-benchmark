#!/usr/bin/env python3
"""officeval_086 的逐题参考解法。

这不是批量造题器：数据映射、两个饼图的几何位置、颜色和连接线均
是根据本题原始 Query、输入表和参考截图逐项确定的。
"""

from __future__ import annotations

import argparse
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from lxml import etree
from openpyxl import load_workbook
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.data_source import (
    AxDataSource,
    NumData,
    NumVal,
    StrData,
    StrRef,
    StrVal,
)
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


OUTPUT_NAME = "数据1_复合饼图_Sheet2.xlsx"
EMU_PER_CM = 360000
XDR_NS = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def add_chart_cache(chart: PieChart, categories: list[str], values: list[float]) -> None:
    """保留单元格引用的同时写入图表缓存，便于离线渲染和检查。"""
    series = chart.ser[0]
    if series.cat.numRef is not None:
        formula = series.cat.numRef.f
    else:
        formula = series.cat.strRef.f
    series.cat = AxDataSource(
        strRef=StrRef(
            f=formula,
            strCache=StrData(
                ptCount=len(categories),
                pt=[StrVal(idx=index, v=value) for index, value in enumerate(categories)],
            ),
        )
    )
    series.val.numRef.numCache = NumData(
        formatCode="0%",
        ptCount=len(values),
        pt=[NumVal(idx=index, v=value) for index, value in enumerate(values)],
    )


def style_pie(
    chart: PieChart,
    categories: list[str],
    values: list[float],
    colors: list[str],
    *,
    width_cm: float,
    height_cm: float,
) -> None:
    chart.width = width_cm
    chart.height = height_cm
    chart.varyColors = True
    chart.legend = None
    chart.title = None
    chart.dataLabels = DataLabelList(
        showLegendKey=False,
        showVal=False,
        showCatName=True,
        showSerName=False,
        showPercent=True,
        showLeaderLines=True,
        separator="\n",
        dLblPos="bestFit",
    )

    points = []
    for index, color in enumerate(colors):
        properties = GraphicalProperties(solidFill=color)
        properties.ln.solidFill = "FFFFFF"
        properties.ln.width = 12700
        properties.ln.prstDash = "solid"
        points.append(DataPoint(idx=index, spPr=properties))
    chart.ser[0].data_points = points
    add_chart_cache(chart, categories, values)


def add_connector_shape(root, *, name: str, shape_id: int, row: int, row_offset_cm: float, flip: bool) -> None:
    """向 Sheet2 加入一条可编辑灰色虚线，连接两个原生饼图。"""
    anchor = etree.Element(f"{{{XDR_NS}}}oneCellAnchor")
    marker = etree.SubElement(anchor, f"{{{XDR_NS}}}from")
    marker_values = {
        "col": "8",
        "colOff": str(int(1.2 * EMU_PER_CM)),
        "row": str(row),
        "rowOff": str(int(row_offset_cm * EMU_PER_CM)),
    }
    for tag, value in marker_values.items():
        etree.SubElement(marker, f"{{{XDR_NS}}}{tag}").text = value

    width = int(4.75 * EMU_PER_CM)
    height = int(0.30 * EMU_PER_CM)
    extent = etree.SubElement(anchor, f"{{{XDR_NS}}}ext")
    extent.set("cx", str(width))
    extent.set("cy", str(height))

    shape = etree.SubElement(anchor, f"{{{XDR_NS}}}sp")
    non_visual = etree.SubElement(shape, f"{{{XDR_NS}}}nvSpPr")
    properties = etree.SubElement(non_visual, f"{{{XDR_NS}}}cNvPr")
    properties.set("id", str(shape_id))
    properties.set("name", name)
    etree.SubElement(non_visual, f"{{{XDR_NS}}}cNvSpPr")

    shape_properties = etree.SubElement(shape, f"{{{XDR_NS}}}spPr")
    transform = etree.SubElement(shape_properties, f"{{{DRAWING_NS}}}xfrm")
    if flip:
        transform.set("flipV", "1")
    offset = etree.SubElement(transform, f"{{{DRAWING_NS}}}off")
    offset.set("x", "0")
    offset.set("y", "0")
    shape_extent = etree.SubElement(transform, f"{{{DRAWING_NS}}}ext")
    shape_extent.set("cx", str(width))
    shape_extent.set("cy", str(height))

    geometry = etree.SubElement(shape_properties, f"{{{DRAWING_NS}}}prstGeom")
    geometry.set("prst", "line")
    etree.SubElement(geometry, f"{{{DRAWING_NS}}}avLst")
    line = etree.SubElement(shape_properties, f"{{{DRAWING_NS}}}ln")
    line.set("w", "12700")
    fill = etree.SubElement(line, f"{{{DRAWING_NS}}}solidFill")
    color = etree.SubElement(fill, f"{{{DRAWING_NS}}}srgbClr")
    color.set("val", "B7B7B7")
    dash = etree.SubElement(line, f"{{{DRAWING_NS}}}prstDash")
    dash.set("val", "dash")
    etree.SubElement(anchor, f"{{{XDR_NS}}}clientData")
    root.append(anchor)


def inject_connectors(path: Path) -> None:
    temporary = path.with_suffix(".tmp.xlsx")
    with ZipFile(path, "r") as source, ZipFile(temporary, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "xl/drawings/drawing1.xml":
                root = etree.fromstring(payload)
                add_connector_shape(root, name="上连接虚线", shape_id=100, row=10, row_offset_cm=0.10, flip=False)
                add_connector_shape(root, name="下连接虚线", shape_id=101, row=18, row_offset_cm=0.05, flip=True)
                payload = etree.tostring(
                    root,
                    xml_declaration=True,
                    encoding="UTF-8",
                    standalone=True,
                )
            target.writestr(info, payload)
    temporary.replace(path)


def solve(input_path: Path, output_path: Path) -> None:
    workbook = load_workbook(input_path)
    source = workbook["Sheet1"]

    # 用 Query 明确给出的业务名称替换输入里的泛化占位名，数值不变。
    labels = {
        "A1": "人员薪酬",
        "A2": "物料",
        "A3": "其他",
        "A5": "其他",
        "B5": "交通",
        "B6": "住宿",
        "B7": "餐饮",
    }
    for cell, value in labels.items():
        source[cell] = value
    for row in range(1, 8):
        for cell in source[row]:
            cell.font = Font(name="微软雅黑", size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center")
    for column in "ABC":
        source.column_dimensions[column].width = 14

    sheet = workbook.create_sheet("Sheet2")
    for column in range(1, 18):
        sheet.column_dimensions[get_column_letter(column)].width = 8.1
    sheet.sheet_view.showGridLines = False

    sheet.merge_cells("B1:P1")
    sheet["B1"] = "费用构成复合饼图"
    sheet["B1"].font = Font(name="微软雅黑", size=18, bold=True, color="1F2937")
    sheet["B1"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("B2:P2")
    sheet["B2"] = "左图展示总体费用结构，右图拆解“其他”费用；数据来源：Sheet1"
    sheet["B2"].font = Font(name="微软雅黑", size=10, color="667085")
    sheet["B2"].alignment = Alignment(horizontal="center")

    main = PieChart()
    main.add_data(Reference(source, min_col=2, min_row=1, max_row=3))
    main.set_categories(Reference(source, min_col=1, min_row=1, max_row=3))
    style_pie(
        main,
        ["人员薪酬", "物料", "其他"],
        [0.55, 0.34, 0.11],
        ["5B9BD5", "ED7D31", "44B3AD"],
        width_cm=15,
        height_cm=12,
    )
    sheet.add_chart(main, "B4")

    detail = PieChart()
    detail.add_data(Reference(source, min_col=3, min_row=5, max_row=7))
    detail.set_categories(Reference(source, min_col=2, min_row=5, max_row=7))
    style_pie(
        detail,
        ["交通", "住宿", "餐饮"],
        [0.15, 0.65, 0.20],
        ["8064A2", "70AD47", "4472C4"],
        width_cm=9.5,
        height_cm=8.5,
    )
    sheet.add_chart(detail, "K8")

    sheet.merge_cells("B29:P29")
    sheet["B29"] = "提示：图表为原生可编辑饼图；在 Sheet1 修改比例后可更新图表。"
    sheet["B29"].font = Font(name="微软雅黑", size=9, color="667085")
    sheet["B29"].alignment = Alignment(horizontal="center")
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_area = "A1:Q30"

    workbook.active = workbook.sheetnames.index("Sheet2")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    inject_connectors(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    solve(args.input, args.output_dir / OUTPUT_NAME)


if __name__ == "__main__":
    main()
