#!/usr/bin/env python3
"""officeval_086 的 Harbor Judge。"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook


TASK_ID = "cowork-evo/officeval-086-editable-composite-pie-chart"
RUBRIC_VERSION = "omegause-officeval-086-adaptation-v0.1"
JUDGE_VERSION = "officeval-086-adaptation-v0.1"
PASS_THRESHOLD = 0.75

TASK_ROOT = Path(__file__).resolve().parents[1]
EXPECTATIONS = json.loads((Path(__file__).with_name("expectations.json")).read_text(encoding="utf-8"))
OUTPUT_NAME = EXPECTATIONS["required_files"][0]

SPECS = {
    "R001": ("指定名称的 XLSX 可正常打开", 2, "hurdle"),
    "R002": ("Sheet1 原数值完整，业务标签映射正确，仅有 Sheet1/Sheet2", 5, "hurdle"),
    "R003": ("Sheet2 含两个与 Sheet1 联动的原生可编辑饼图，不是截图", 4, "hurdle"),
    "R004": ("主饼图与拆分饼图的核心布局、分类和数据比例正确", 5, "positive"),
    "R005": ("左侧主饼图显示人员薪酬 55%、物料 34%、其他 11%", 1, "positive"),
    "R006": ("右侧拆分饼图显示交通 15%、住宿 65%、餐饮 20%", 1, "positive"),
    "R007": ("六个扇区使用可区分配色和白色分隔线", 1, "positive"),
    "R008": ("两个饼图之间有两条灰色虚线连接", 1, "positive"),
    "R009": ("主饼图、拆分饼图的位置、大小比和间距正确", 1, "positive"),
    "R010": ("图例不遮挡，页面标题有意义且无默认“图表标题”", 2, "positive"),
    "R011": ("工作簿可实际渲染为两页 PDF，第二页包含六个分类和标题", 3, "positive"),
    "R012": ("以截图代替可编辑图表，或缺失任一原生饼图", -8, "penalty"),
    "R013": ("修改了任一原始比例数值或丢失 Sheet1 数据", -6, "penalty"),
    "R014": ("新增无关工作表或保留默认“图表标题”", -2, "penalty"),
}


def load_upstream_module():
    path = Path(__file__).with_name("upstream_verifier.py")
    spec = importlib.util.spec_from_file_location("officeval_086_upstream", path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


UPSTREAM = load_upstream_module()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def close(a: float, b: float, tolerance: float = 1e-9) -> bool:
    return abs(float(a) - float(b)) <= tolerance


def result_item(identifier: str, passed: bool, evidence: str) -> dict:
    description, weight, kind = SPECS[identifier]
    return {
        "criterion_id": identifier,
        "description": description,
        "weight": weight,
        "type": kind,
        "passed": bool(passed),
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
    }


def source_check(workbook) -> tuple[bool, str, bool, bool]:
    sheet_names_ok = workbook.sheetnames == EXPECTATIONS["sheet_names"]
    if "Sheet1" not in workbook.sheetnames:
        return False, "缺少 Sheet1", True, not sheet_names_ok
    sheet = workbook["Sheet1"]
    main_labels = [sheet[cell].value for cell in ("A1", "A2", "A3")]
    main_values = [sheet[cell].value for cell in ("B1", "B2", "B3")]
    detail_labels = [sheet[cell].value for cell in ("B5", "B6", "B7")]
    detail_values = [sheet[cell].value for cell in ("C5", "C6", "C7")]
    labels_ok = main_labels == EXPECTATIONS["main_labels"] and detail_labels == EXPECTATIONS["detail_labels"]
    values_ok = (
        all(close(actual, expected) for actual, expected in zip(main_values, EXPECTATIONS["main_values"]))
        and all(close(actual, expected) for actual, expected in zip(detail_values, EXPECTATIONS["detail_values"]))
    )
    merge_ok = "A5:A7" in {str(item) for item in sheet.merged_cells.ranges}
    passed = sheet_names_ok and labels_ok and values_ok and merge_ok
    evidence = (
        f"工作表={workbook.sheetnames}，主分类={list(zip(main_labels, main_values))}，"
        f"明细={list(zip(detail_labels, detail_values))}，合并 A5:A7={merge_ok}"
    )
    severe_loss = not values_ok or "Sheet1" not in workbook.sheetnames
    extra_sheets = not sheet_names_ok
    return passed, evidence, severe_loss, extra_sheets


def native_chart_check(workbook, model) -> tuple[bool, str, bool]:
    if "Sheet2" not in workbook.sheetnames:
        return False, "缺少 Sheet2", True
    sheet = workbook["Sheet2"]
    charts = sheet._charts
    image_count = len(sheet._images)
    types = [type(chart).__name__ for chart in charts]
    formulas = []
    for chart in charts:
        for series in chart.ser:
            value_formula = series.val.numRef.f if series.val is not None and series.val.numRef is not None else ""
            category_formula = ""
            if series.cat is not None:
                if series.cat.strRef is not None:
                    category_formula = series.cat.strRef.f
                elif series.cat.numRef is not None:
                    category_formula = series.cat.numRef.f
            formulas.append((category_formula, value_formula))
    linked = len(formulas) == 2 and all("Sheet1" in cat and "Sheet1" in val for cat, val in formulas)
    model_types = [chart.chart_types for chart in model.charts]
    native_ok = len(charts) == EXPECTATIONS["chart_count"] and all(name == "PieChart" for name in types)
    passed = native_ok and len(model.charts) == 2 and linked and image_count == 0
    flattened = not native_ok or image_count > 0
    return passed, f"图表={types}，OOXML={model_types}，引用={formulas}，图片={image_count}", flattened


def render_check(path: Path) -> tuple[bool, str]:
    executables = [shutil.which(name) for name in ("libreoffice", "pdfinfo", "pdftotext")]
    if not all(executables):
        return False, "缺少 libreoffice/pdfinfo/pdftotext"
    libreoffice, pdfinfo, pdftotext = executables
    with tempfile.TemporaryDirectory(prefix="cowork-evo-086-") as temporary:
        root = Path(temporary)
        environment = os.environ.copy()
        for key in ("HTTP_PROXY", "HTTPS_PROXY", "ALL_PROXY", "http_proxy", "https_proxy", "all_proxy"):
            environment.pop(key, None)
        completed = subprocess.run(
            [
                libreoffice,
                f"-env:UserInstallation={(root / 'lo-profile').as_uri()}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(root),
                str(path),
            ],
            capture_output=True,
            text=True,
            timeout=180,
            check=False,
            env=environment,
        )
        pdf = root / f"{path.stem}.pdf"
        if completed.returncode != 0 or not pdf.is_file():
            return False, (completed.stderr or completed.stdout or "转换失败")[-1000:]
        info = subprocess.run([pdfinfo, str(pdf)], capture_output=True, text=True, timeout=30, check=False)
        text = subprocess.run([pdftotext, "-layout", str(pdf), "-"], capture_output=True, text=True, timeout=30, check=False)
        page_match = re.search(r"^Pages:\s+(\d+)", info.stdout, re.MULTILINE)
        pages = int(page_match.group(1)) if page_match else 0
        normalized = re.sub(r"\s+", "", text.stdout)
        required = [EXPECTATIONS["worksheet_title"], *EXPECTATIONS["main_labels"], *EXPECTATIONS["detail_labels"]]
        missing = [value for value in required if value not in normalized]
        passed = info.returncode == 0 and text.returncode == 0 and pages == EXPECTATIONS["render_pages"] and not missing
        return passed, f"页数={pages}，缺失={missing}，PDF={pdf.stat().st_size} bytes"


def evaluate(artifact_dir: Path) -> dict:
    output = artifact_dir / OUTPUT_NAME
    results: list[dict] = []
    if not output.is_file():
        results.append(result_item("R001", False, f"缺少 {OUTPUT_NAME}"))
        return build_result(results, output)
    try:
        workbook = load_workbook(output, data_only=False)
        value_workbook = load_workbook(output, data_only=True)
        model = UPSTREAM.load_model(output)
    except Exception as exc:
        results.append(result_item("R001", False, f"无法打开：{type(exc).__name__}: {exc}"))
        return build_result(results, output)
    results.append(result_item("R001", True, f"XLSX 可打开，工作表={workbook.sheetnames}"))

    source_ok, source_evidence, severe_loss, extra_sheets = source_check(workbook)
    results.append(result_item("R002", source_ok, source_evidence))
    native_ok, native_evidence, flattened = native_chart_check(workbook, model)
    results.append(result_item("R003", native_ok, native_evidence))

    hits = UPSTREAM.score_dimension2(model, workbook, value_workbook)
    upstream_map = {
        "R004": hits[0],
        "R005": hits[1],
        "R006": hits[2],
        "R007": hits[3],
        "R008": hits[4],
        "R009": hits[5],
    }
    for identifier, hit in upstream_map.items():
        results.append(result_item(identifier, hit.passed, hit.detail))
    legend, title = hits[6], hits[7]
    page_title_ok = "Sheet2" in workbook.sheetnames and workbook["Sheet2"]["B1"].value == EXPECTATIONS["worksheet_title"]
    results.append(
        result_item(
            "R010",
            legend.passed and title.passed and page_title_ok,
            f"{legend.detail}；{title.detail}；页面标题={page_title_ok}",
        )
    )
    rendered, render_evidence = render_check(output)
    results.append(result_item("R011", rendered, render_evidence))

    results.append(result_item("R012", flattened, native_evidence))
    results.append(result_item("R013", severe_loss, source_evidence))
    default_title = any("图表标题" in (chart.title or "") for chart in model.charts)
    results.append(result_item("R014", extra_sheets or default_title, f"额外工作表={extra_sheets}，默认标题={default_title}"))
    return build_result(results, output)


def build_result(results: list[dict], artifact: Path) -> dict:
    present = {item["criterion_id"] for item in results}
    for identifier in SPECS:
        if identifier not in present:
            results.append(result_item(identifier, False, "因门槛失败未执行"))
    results.sort(key=lambda item: item["criterion_id"])
    positive_total = sum(item["weight"] for item in results if item["weight"] > 0)
    earned = sum(item["weight"] for item in results if item["weight"] > 0 and item["passed"])
    penalties = sum(item["weight"] for item in results if item["weight"] < 0 and item["passed"])
    reward = max(0.0, min(1.0, (earned + penalties) / positive_total))
    hurdle_failures = [item["criterion_id"] for item in results if item["type"] == "hurdle" and not item["passed"]]
    if hurdle_failures:
        reward = 0.0
    normalized = []
    for item in results:
        score = (0.0 if item["passed"] else 1.0) if item["type"] == "penalty" else item["score"]
        normalized.append(
            {
                "criterion_id": item["criterion_id"],
                "score": score,
                "evidence": item["evidence"],
                "method": item["method"],
                "raw": {
                    "description": item["description"],
                    "weight": item["weight"],
                    "type": item["type"],
                    "passed": item["passed"],
                },
            }
        )
    digest = sha256_file(artifact) if artifact.is_file() else hashlib.sha256(b"").hexdigest()
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [digest],
        "criterion_results": normalized,
        "reward": round(reward, 6),
        "passed": reward >= PASS_THRESHOLD,
        "hurdle_failures": hurdle_failures,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path)
    parser.add_argument("--reward-file", type=Path)
    args = parser.parse_args()
    result = evaluate(args.artifact_dir)
    rendered = json.dumps(result, ensure_ascii=False, indent=2)
    print(rendered)
    if args.result:
        args.result.parent.mkdir(parents=True, exist_ok=True)
        args.result.write_text(rendered + "\n", encoding="utf-8")
    if args.reward_file:
        args.reward_file.parent.mkdir(parents=True, exist_ok=True)
        args.reward_file.write_text(f"{result['reward']:.6f}\n", encoding="utf-8")
    raise SystemExit(0 if result["passed"] else 1)


if __name__ == "__main__":
    main()
