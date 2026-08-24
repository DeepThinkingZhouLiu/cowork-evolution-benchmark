#!/usr/bin/env python3
"""officeval_087 的逐题参考解法。

本文件只处理上游给定的两个固定工作簿。题干、拼音、作业内容和图片均来自
“识字主题分层作业.xlsx”；目标样式、合并区域和空白工作表均来自
“分层作业设计.xlsx”。它不是可以替换输入来批量造题的通用生成器。
"""

from __future__ import annotations

import argparse
import io
import os
import re
import tempfile
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Font
from openpyxl.utils.units import cm_to_EMU


TEMPLATE_NAME = "分层作业设计.xlsx"
SOURCE_NAME = "识字主题分层作业.xlsx"
OUTPUT_NAME = "分层作业设计_填充完成.xlsx"
EXPECTED_SHEETS = ["Sheet1", "Sheet2", "Sheet3"]
EXPECTED_MERGES = {
    "A1:F1",
    "B3:F3",
    "C4:F4",
    "C5:F5",
    "C6:F6",
    "A4:A6",
    "A7:A9",
    "B7:F9",
}

# 以下三段内容逐项来自源工作簿 Sheet1 的“基础类、提升类、拓展类”区域。
BASE_TEXT = """1.给词语选择正确的读音画“√”。
晨露：chén lù（ ） chēn lù（ ）
清风：qīng fēng（ ） qīn fēng（ ）
细雨：xì yǔ（ ） xī yǔ（ ）
晚霞：wǎn xiá（ ） xiá wǎn（ ）

2.看拼音，写词语。
qīng chén：____　xiǎo yǔ：____　hé fēng：____
太阳从山后慢慢升起，小雨轻轻落在窗前，和风送来阵阵清香。
huā duǒ：____　niǎo ér：____　fēi：____
花园里的花开得正好，鸟儿从树梢掠过，飞进了安静的小院。





3.你能圈出这几个词中表示自然景物的字吗？
晨露　细雨　晚霞

4.连一连。
晨露　午阳　晚霞　夜风
亮　暖　红　轻"""

IMPROVE_TEXT = """1.词语接龙。
清晨→（　）→（　）→（　）

2.照样子，看图说一说。
例：微风轻　小鸟（　）　柳叶（　）　溪水（　）
纸鸢（　）　云朵（　）






3.仔细观察图片，填一填。
清风吹过，河岸边铺满了嫩嫩的____，草地上点缀着____。
五颜六色的____，清亮的____中，____慢慢游来游去，
还轻轻划出了____纹。远处的____上开着浅浅的____。
一群小鸟正快活地____林间。
田字格补写：____、____、____、____、____、____"""

EXPAND_TEXT = """1.和家人到附近公园或校园走一走，记录三种自然景物，回到班级分享。

2.搜集两个描写季节的词语，和同学交流。"""

FEEDBACK_TEXT = """自我评价：今天我完成了_______颗星，我已经累计了_______颗星，请继续保持！
教师评价："""


def set_font(cell, size: float, *, bold: bool = False) -> None:
    cell.font = Font(name="宋体", size=size, bold=bold, color="FF000000")


def set_text_cell(
    cell,
    text: str,
    *,
    size: float = 11,
    bold: bool = False,
    horizontal: str | None = "left",
    vertical: str = "top",
    wrap: bool = True,
) -> None:
    cell.value = text
    set_font(cell, size, bold=bold)
    cell.alignment = Alignment(
        horizontal=horizontal,
        vertical=vertical,
        wrap_text=wrap,
    )


def source_media(source_path: Path, name: str) -> bytes:
    with zipfile.ZipFile(source_path) as archive:
        return archive.read(f"xl/media/{name}")


def add_source_image(
    worksheet,
    payload: bytes,
    *,
    column: int,
    row: int,
    width_cm: float,
    height_cm: float,
    column_offset_cm: float = 0,
    row_offset_cm: float = 0,
) -> None:
    """把源工作簿中的真实图片放入目标合并区域；行列参数从 1 开始。"""

    image = Image(io.BytesIO(payload))
    image.width = width_cm / 2.54 * 96
    image.height = height_cm / 2.54 * 96
    marker = AnchorMarker(
        col=column - 1,
        colOff=cm_to_EMU(column_offset_cm),
        row=row - 1,
        rowOff=cm_to_EMU(row_offset_cm),
    )
    image.anchor = OneCellAnchor(
        _from=marker,
        ext=XDRPositiveSize2D(
            cx=cm_to_EMU(width_cm),
            cy=cm_to_EMU(height_cm),
        ),
    )
    worksheet.add_image(image)


def validate_inputs(template_path: Path, source_path: Path) -> None:
    template = load_workbook(template_path, read_only=False, data_only=False)
    if template.sheetnames != EXPECTED_SHEETS:
        raise ValueError(f"模板工作表不符：{template.sheetnames!r}")
    if {str(item) for item in template["Sheet1"].merged_cells.ranges} != EXPECTED_MERGES:
        raise ValueError("模板的 8 个合并区域与本题预期不符")
    for name in ("Sheet2", "Sheet3"):
        worksheet = template[name]
        if any(cell.value not in (None, "") for row in worksheet.iter_rows() for cell in row):
            raise ValueError(f"模板 {name} 不再为空白")
    source = load_workbook(source_path, read_only=False, data_only=False)
    if source.sheetnames != ["Sheet1"]:
        raise ValueError(f"源工作簿工作表不符：{source.sheetnames!r}")
    source_text = "\n".join(
        str(cell.value)
        for row in source["Sheet1"].iter_rows()
        for cell in row
        if cell.value not in (None, "")
    )
    for required in (
        "识字主题《四季小景》分层作业",
        "给词语选择正确的读音",
        "词语接龙",
        "和家人到附近公园或校园走一走",
    ):
        if required not in source_text:
            raise ValueError(f"源工作簿缺少本题真实内容：{required}")


def normalize_drawing_prefixes(path: Path) -> None:
    """让 drawing XML 使用显式 xdr 前缀，保持与上游 Verifier 的解析约定一致。

    openpyxl 默认把 spreadsheetDrawing 声明为无前缀默认命名空间；上游
    Verifier 按 Office 原文件常见的 ``xdr:`` 前缀读取同一命名空间。这里只改变
    等价的 XML 前缀，不改变图片、锚点、尺寸或工作簿内容。
    """

    drawing_namespace = (
        "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
    )
    drawing_tags = (
        "oneCellAnchor",
        "twoCellAnchor",
        "absoluteAnchor",
        "from",
        "to",
        "col",
        "colOff",
        "row",
        "rowOff",
        "ext",
        "pos",
        "pic",
        "nvPicPr",
        "cNvPr",
        "cNvPicPr",
        "blipFill",
        "spPr",
        "clientData",
    )
    with tempfile.NamedTemporaryFile(
        prefix=f"{path.stem}-",
        suffix=".xlsx",
        dir=path.parent,
        delete=False,
    ) as handle:
        temporary_path = Path(handle.name)
    try:
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
            temporary_path, "w", compression=zipfile.ZIP_DEFLATED
        ) as target:
            for info in source.infolist():
                payload = source.read(info.filename)
                if info.filename.startswith("xl/drawings/drawing") and info.filename.endswith(".xml"):
                    xml = payload.decode("utf-8")
                    xml = xml.replace(
                        f'<wsDr xmlns="{drawing_namespace}">',
                        f'<xdr:wsDr xmlns:xdr="{drawing_namespace}">',
                    ).replace("</wsDr>", "</xdr:wsDr>")
                    for tag in drawing_tags:
                        xml = re.sub(
                            rf"(<\/?){tag}(?=[ >])",
                            rf"\1xdr:{tag}",
                            xml,
                        )
                    payload = xml.encode("utf-8")
                target.writestr(info, payload)
        os.replace(temporary_path, path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def fill_workbook(template_path: Path, source_path: Path, output_path: Path) -> None:
    validate_inputs(template_path, source_path)
    workbook = load_workbook(template_path, read_only=False, data_only=False)
    worksheet = workbook["Sheet1"]

    # 上游细则给出的列宽是 Excel 显示值，逐列明确设置；合并和边框继续沿用模板。
    for column, width in {
        "A": 12.75,
        "B": 14.63,
        "C": 13.13,
        "D": 17.38,
        "E": 12.25,
        "F": 15.50,
    }.items():
        worksheet.column_dimensions[column].width = width

    set_text_cell(
        worksheet["A1"],
        "识字主题《四季小景》分层作业",
        size=18,
        bold=True,
        horizontal="center",
        vertical="center",
    )
    worksheet.row_dimensions[1].height = 39

    for ref, text in (("A2", "班级"), ("C2", "任教教师"), ("E2", "时间")):
        set_text_cell(
            worksheet[ref],
            text,
            size=11,
            bold=True,
            horizontal="center",
            vertical="center",
        )
    for ref in ("B2", "D2", "F2"):
        worksheet[ref].value = None

    set_text_cell(
        worksheet["A3"],
        "作业设计简介（学习目标、题型说明等）",
        bold=True,
        horizontal="center",
        vertical="center",
    )
    set_text_cell(
        worksheet["B3"],
        "识字主题《四季小景》分层作业；基础类、提升类为必做题，拓展类为选做题。",
        bold=False,
        horizontal="left",
        vertical="center",
    )
    worksheet.row_dimensions[3].height = 68

    set_text_cell(
        worksheet["A4"],
        "作业\n设计\n内容",
        size=14,
        bold=True,
        horizontal="center",
        vertical="center",
    )
    for ref, text in (
        ("B4", "基础层作业设计"),
        ("B5", "提高层作业设计"),
        ("B6", "拓展层作业设计"),
    ):
        set_text_cell(
            worksheet[ref],
            text,
            bold=True,
            horizontal="center",
            vertical="center",
        )

    set_text_cell(worksheet["C4"], BASE_TEXT, size=10.5, vertical="top")
    set_text_cell(worksheet["C5"], IMPROVE_TEXT, size=10.5, vertical="top")
    set_text_cell(worksheet["C6"], EXPAND_TEXT, size=11, vertical="center")
    worksheet.row_dimensions[4].height = 330
    worksheet.row_dimensions[5].height = 420
    worksheet.row_dimensions[6].height = 145

    set_text_cell(
        worksheet["A7"],
        "作业情况反馈",
        bold=True,
        horizontal="center",
        vertical="center",
    )
    set_text_cell(
        worksheet["B7"],
        FEEDBACK_TEXT,
        size=11,
        horizontal="left",
        vertical="center",
    )

    # 两张风景插图和田字格均直接复用源工作簿的媒体，不重画、不生成。
    add_source_image(
        worksheet,
        source_media(source_path, "image7.jpeg"),
        column=3,
        row=5,
        width_cm=4.5,
        height_cm=2.25,
        column_offset_cm=0.15,
        row_offset_cm=3.0,
    )
    add_source_image(
        worksheet,
        source_media(source_path, "image21.png"),
        column=5,
        row=5,
        width_cm=5.6,
        height_cm=2.8,
        column_offset_cm=0,
        row_offset_cm=9.2,
    )

    # 基础层的 6 个田字格来自源表 image1-image6，放在第2题与第3题之间的留白行。
    for index, media_name in enumerate(
        ("image1.png", "image2.png", "image3.png", "image4.png", "image5.png", "image6.png")
    ):
        add_source_image(
            worksheet,
            source_media(source_path, media_name),
            column=3,
            row=4,
            width_cm=0.95 if media_name != "image4.png" else 1.8,
            height_cm=0.95,
            column_offset_cm=0.10 + index * 1.75,
            row_offset_cm=6.35,
        )

    worksheet.sheet_view.showGridLines = False
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.print_area = "A1:F9"
    worksheet.freeze_panes = None

    if workbook.sheetnames != EXPECTED_SHEETS:
        raise ValueError("保存前工作表名称、数量或顺序发生变化")
    if {str(item) for item in worksheet.merged_cells.ranges} != EXPECTED_MERGES:
        raise ValueError("保存前模板合并区域发生变化")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    normalize_drawing_prefixes(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=Path("/app/data/input_files"))
    parser.add_argument("--output-dir", type=Path, default=Path("/app/output"))
    args = parser.parse_args()
    fill_workbook(
        args.input_dir / TEMPLATE_NAME,
        args.input_dir / SOURCE_NAME,
        args.output_dir / OUTPUT_NAME,
    )


if __name__ == "__main__":
    main()
