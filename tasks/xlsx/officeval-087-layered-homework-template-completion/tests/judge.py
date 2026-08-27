#!/usr/bin/env python3
"""officeval_087 的 Harbor Judge。

上游 Verifier 负责原始细则，本 Judge 额外核对真实源文本、选定源媒体的 SHA256、
模板边框与填充、空白工作表、可编辑性和 LibreOffice 实际渲染。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import subprocess
import tempfile
import warnings
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

import upstream_verifier


OUTPUT_NAME = "分层作业设计_填充完成.xlsx"
TEMPLATE_NAME = "分层作业设计.xlsx"
TASK_ID = "cowork-evo/officeval-087-layered-homework-template-completion"
RUBRIC_VERSION = "omegause-officeval-087-adaptation-v0.1"
JUDGE_VERSION = "officeval-087-adaptation-v0.1"
EXPECTED_SHEETS = ["Sheet1", "Sheet2", "Sheet3"]
EXPECTED_MERGES = {
    "A1:F1",
    "B3:F3",
    "C4:F4",
    "C5:F5",
    "C6:F6",
    "A4:A6",
    "A7:A9",
    "B7:F9",
}

CRITERION_SPECS = {
    "R001": ("指定 XLSX 可打开，且工作表名称、数量和顺序正确", 2, "hurdle"),
    "R002": ("识字源表的标题、三层作业、田字格语境和评价内容完整迁入", 8, "positive"),
    "R003": ("模板 8 个合并区域、基本信息空白区及 Sheet2/Sheet3 完整保留", 4, "positive"),
    "R004": ("指定列宽、行高、字体、对齐和自动换行满足模板细则", 4, "positive"),
    "R005": ("8 个真实源媒体对象原样复用，且两张山水图位置和尺寸正确", 4, "positive"),
    "R006": ("模板表格边框、填充和固定标签未被无关改动", 3, "positive"),
    "R007": ("LibreOffice 实际渲染后关键内容可读且没有横向分页破坏", 3, "positive"),
    "R008": ("三层作业仍保存在可编辑单元格中，未被整表截图替代", 2, "positive"),
    "R009": ("两处及以上真实源内容缺失、改写或错层", -8, "penalty"),
    "R010": ("模板结构、空白工作表、边框或固定标签被破坏", -8, "penalty"),
    "R011": ("关键版式、图片位置或实际渲染存在明显缺陷", -5, "penalty"),
}


def criterion(identifier: str, passed: bool, evidence: str = "") -> dict:
    description, weight, kind = CRITERION_SPECS[identifier]
    return {
        "criterion_id": identifier,
        "description": description,
        "weight": weight,
        "type": kind,
        "passed": bool(passed),
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
    }


def compact(value) -> str:
    return re.sub(r"\s+", "", "" if value is None else str(value))


def contains_all(value, required: list[str]) -> list[str]:
    actual = compact(value)
    return [item for item in required if compact(item) not in actual]


def is_blank_sheet(worksheet) -> bool:
    has_values = any(
        cell.value not in (None, "")
        for row in worksheet.iter_rows()
        for cell in row
        if not isinstance(cell, MergedCell)
    )
    return (
        not has_values
        and not worksheet.merged_cells.ranges
        and not getattr(worksheet, "_images", [])
        and not getattr(worksheet, "_charts", [])
    )


def content_errors(worksheet, expected: dict) -> list[str]:
    errors = []
    if compact(worksheet["A1"].value) != compact(expected["title"]):
        errors.append("标题不匹配")
    if compact(worksheet["B3"].value) != compact(expected["intro"]):
        errors.append("作业简介不匹配")
    for ref, key, label in (
        ("C4", "base_required", "基础层"),
        ("C5", "improve_required", "提高层"),
        ("C6", "expand_required", "拓展层"),
        ("B7", "feedback_required", "反馈区"),
    ):
        missing = contains_all(worksheet[ref].value, expected[key])
        if missing:
            errors.append(f"{label}缺失={missing}")
    base_slots = len(re.findall(r"_{4,}", str(worksheet["C4"].value or "")))
    improve_slots = len(re.findall(r"_{4,}", str(worksheet["C5"].value or "")))
    if base_slots != 6:
        errors.append(f"基础层拼音填空槽={base_slots}，期望=6")
    if improve_slots != 15:
        errors.append(f"提高层田字格填空槽={improve_slots}，期望=15")
    if "\n" not in str(worksheet["B7"].value or ""):
        errors.append("教师评价未另起一行")
    return errors


def structure_errors(workbook) -> list[str]:
    errors = []
    worksheet = workbook["Sheet1"]
    actual_merges = {str(item) for item in worksheet.merged_cells.ranges}
    if actual_merges != EXPECTED_MERGES:
        errors.append(
            f"合并区域改变：缺失={sorted(EXPECTED_MERGES - actual_merges)}，"
            f"新增={sorted(actual_merges - EXPECTED_MERGES)}"
        )
    expected_labels = {
        "A2": "班级",
        "C2": "任教教师",
        "E2": "时间",
        "A3": "作业设计简介（学习目标、题型说明等）",
        "A4": "作业设计内容",
        "B4": "基础层作业设计",
        "B5": "提高层作业设计",
        "B6": "拓展层作业设计",
        "A7": "作业情况反馈",
    }
    for ref, expected in expected_labels.items():
        if compact(worksheet[ref].value) != compact(expected):
            errors.append(f"固定标签 {ref} 改变")
    for ref in ("B2", "D2", "F2"):
        if worksheet[ref].value not in (None, ""):
            errors.append(f"基本信息填写区 {ref} 不再空白")
    for name in ("Sheet2", "Sheet3"):
        if not is_blank_sheet(workbook[name]):
            errors.append(f"{name} 不再保持空白")
    return errors


def layout_errors(workbook, output_path: Path) -> list[str]:
    worksheet = workbook["Sheet1"]
    errors = []
    widths = {"A": 12.75, "B": 14.63, "C": 13.13, "D": 17.38, "E": 12.25, "F": 15.5}
    for column, expected in widths.items():
        actual = worksheet.column_dimensions[column].width
        if actual is None or abs(float(actual) - expected) > 0.05:
            errors.append(f"{column}列宽={actual}，期望={expected}")
    heights = {1: (38.9, 39.1), 3: (67.9, 68.1), 4: (250, 350), 5: (350, 450)}
    for row, (low, high) in heights.items():
        actual = worksheet.row_dimensions[row].height
        if actual is None or not low <= float(actual) <= high:
            errors.append(f"第{row}行高={actual}，期望范围={low}-{high}")
    style_checks = {
        "A1宋体18加粗居中": (
            worksheet["A1"].font.name == "宋体"
            and worksheet["A1"].font.sz is not None
            and abs(float(worksheet["A1"].font.sz) - 18) < 0.2
            and bool(worksheet["A1"].font.bold)
            and worksheet["A1"].alignment.horizontal == "center"
            and worksheet["A1"].alignment.vertical == "center"
        ),
        "A3宋体11加粗换行居中": (
            worksheet["A3"].font.name == "宋体"
            and worksheet["A3"].font.sz is not None
            and abs(float(worksheet["A3"].font.sz) - 11) < 0.2
            and bool(worksheet["A3"].font.bold)
            and bool(worksheet["A3"].alignment.wrap_text)
            and worksheet["A3"].alignment.horizontal == "center"
        ),
        "三层内容自动换行": all(
            bool(worksheet[ref].alignment.wrap_text) for ref in ("C4", "C5", "C6")
        ),
        "反馈区宋体11自动换行": (
            worksheet["B7"].font.name == "宋体"
            and worksheet["B7"].font.sz is not None
            and abs(float(worksheet["B7"].font.sz) - 11) < 0.2
            and bool(worksheet["B7"].alignment.wrap_text)
        ),
    }
    errors.extend(name for name, passed in style_checks.items() if not passed)

    # 保留上游对整套逐点样式细则的独立检查，避免本地 Judge 漏掉某个模板点。
    upstream_results = upstream_verifier.evaluate_dimension2(workbook, output_path)
    layout_names = {
        "列宽：A-F列宽固定符合细则",
        "标题区域 A1:F1 格式与文本",
        "基本信息行 A2/C2/E2 字段格式且 B2/D2/F2 空白",
        "A3 作业简介标题",
        "B3:F3 作业简介内容",
        "A4:A6 层级总标题",
        "B4 基础层标题",
        "B5 提高层标题",
        "B6 拓展层标题",
        "A7:A9 反馈标题区域",
        "B7:F9 反馈内容区域",
        "Sheet1 8个合并单元格区域完全一致且无新增/取消",
    }
    for item in upstream_results:
        if item.name in layout_names and not item.passed:
            errors.append(f"上游细则未通过：{item.name}")
    return sorted(set(errors))


def media_hashes(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as archive:
        return [
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("xl/media/") and not name.endswith("/")
        ]


def media_errors(workbook, output_path: Path, expected: dict) -> list[str]:
    worksheet = workbook["Sheet1"]
    errors = []
    actual_hashes = media_hashes(output_path)
    if Counter(actual_hashes) != Counter(expected["selected_source_media_sha256"]):
        errors.append("输出媒体并非选定的 8 个真实源图片，或数量发生变化")
    if len(getattr(worksheet, "_images", [])) != 8:
        errors.append(f"Sheet1 图片数={len(getattr(worksheet, '_images', []))}，期望=8")
    specs = upstream_verifier.get_drawing_image_specs(output_path, 1)
    base_specs = [item for item in specs if item.from_row == 3]
    improve_specs = [item for item in specs if item.from_row == 4]
    if len(base_specs) != 6:
        errors.append(f"基础层真实田字格图片锚点数={len(base_specs)}，期望=6")
    if len(improve_specs) != 2:
        errors.append(f"提高层山水图片锚点数={len(improve_specs)}，期望=2")
    landscape = [
        item
        for item in improve_specs
        if 5.5 <= item.width_cm <= 5.7 and 2.7 <= item.height_cm <= 2.9
    ]
    if len(landscape) != 1:
        errors.append(f"5.5–5.7cm × 2.7–2.9cm 山水图片数={len(landscape)}，期望=1")
    elif not upstream_verifier.c5_landscape_not_covering_text(worksheet, landscape[0]):
        errors.append("第3题山水图片越界或落入主要文字区")
    return errors


def color_signature(color) -> tuple | None:
    if color is None:
        return None
    return (
        color.type,
        color.rgb,
        color.indexed,
        color.auto,
        color.theme,
        color.tint,
    )


def border_signature(border) -> tuple:
    sides = []
    for name in ("left", "right", "top", "bottom", "diagonal", "vertical", "horizontal"):
        side = getattr(border, name)
        sides.append(
            (name, None, None)
            if side is None
            else (name, side.style, color_signature(side.color))
        )
    return (
        tuple(sides),
        border.outline,
        border.diagonalUp,
        border.diagonalDown,
    )


def fill_signature(fill) -> tuple:
    return (
        fill.fill_type,
        fill.patternType,
        color_signature(fill.fgColor),
        color_signature(fill.bgColor),
    )


def preservation_errors(template, candidate) -> list[str]:
    before = template["Sheet1"]
    after = candidate["Sheet1"]
    errors = []
    for row in range(1, 10):
        for column in range(1, 7):
            left = before.cell(row, column)
            right = after.cell(row, column)
            if border_signature(left.border) != border_signature(right.border):
                errors.append(f"{right.coordinate} 边框改变")
            if fill_signature(left.fill) != fill_signature(right.fill):
                errors.append(f"{right.coordinate} 填充改变")
    return errors


def render_check(path: Path, expected: dict) -> tuple[bool, str]:
    libreoffice = shutil.which("libreoffice")
    pdftotext = shutil.which("pdftotext")
    pdfinfo = shutil.which("pdfinfo")
    if libreoffice is None or pdftotext is None or pdfinfo is None:
        return False, "缺少 libreoffice、pdftotext 或 pdfinfo"
    with tempfile.TemporaryDirectory(prefix="cowork-evo-087-") as temp_dir:
        temp_root = Path(temp_dir)
        profile_uri = (temp_root / "lo-profile").resolve().as_uri()
        converted = subprocess.run(
            [
                libreoffice,
                "--headless",
                f"-env:UserInstallation={profile_uri}",
                "--convert-to",
                "pdf",
                "--outdir",
                temp_dir,
                str(path),
            ],
            capture_output=True,
            text=True,
            timeout=300,
            check=False,
        )
        pdf_path = temp_root / f"{path.stem}.pdf"
        if converted.returncode != 0 or not pdf_path.is_file():
            return False, (converted.stderr or converted.stdout or "PDF 转换失败")[-1000:]
        info = subprocess.run(
            [pdfinfo, str(pdf_path)],
            capture_output=True,
            text=True,
            timeout=120,
            check=False,
        )
        extracted = subprocess.run(
            [pdftotext, "-raw", str(pdf_path), "-"],
            capture_output=True,
            text=True,
            timeout=120,
            check=False,
        )
        if info.returncode != 0 or extracted.returncode != 0:
            return False, (info.stderr or extracted.stderr or "PDF 检查失败")[-1000:]
        match = re.search(r"^Pages:\s+(\d+)", info.stdout, flags=re.MULTILINE)
        pages = int(match.group(1)) if match else 0
        rendered = compact(extracted.stdout)
    required = [
        expected["title"],
        "给词语选择正确的读音",
        "词语接龙",
        "和家人到附近公园或校园走一走",
        "自我评价",
        "教师评价",
    ]
    missing = [item for item in required if compact(item) not in rendered]
    page_ok = 2 <= pages <= 4
    return page_ok and not missing, f"页数={pages}，渲染缺失={missing}"


def evaluate(artifact_dir: Path) -> dict:
    output_path = artifact_dir / OUTPUT_NAME
    results = []
    if not output_path.is_file():
        results.append(criterion("R001", False, f"缺少 {OUTPUT_NAME}"))
        return build_result(results, output_path)

    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            workbook = load_workbook(output_path, data_only=False)
    except Exception as exc:
        results.append(criterion("R001", False, f"文件无法打开：{exc}"))
        return build_result(results, output_path)

    sheets_ok = workbook.sheetnames == EXPECTED_SHEETS
    results.append(criterion("R001", sheets_ok, f"工作表={workbook.sheetnames}"))
    if not sheets_ok:
        return build_result(results, output_path)

    expected_path = Path(__file__).with_name("expected-content.json")
    expected = json.loads(expected_path.read_text(encoding="utf-8"))
    worksheet = workbook["Sheet1"]

    content_issues = content_errors(worksheet, expected)
    results.append(
        criterion(
            "R002",
            not content_issues,
            "真实源内容和 6/15 个填空槽完整" if not content_issues else "；".join(content_issues[:20]),
        )
    )

    structure_issues = structure_errors(workbook)
    results.append(
        criterion(
            "R003",
            not structure_issues,
            "8 个合并区域、基本信息空白区和两个空白 Sheet 完整"
            if not structure_issues
            else "；".join(structure_issues[:20]),
        )
    )

    layout_issues = layout_errors(workbook, output_path)
    results.append(
        criterion(
            "R004",
            not layout_issues,
            "列宽、行高、字体、对齐和换行符合细则" if not layout_issues else "；".join(layout_issues[:20]),
        )
    )

    image_issues = media_errors(workbook, output_path, expected)
    results.append(
        criterion(
            "R005",
            not image_issues,
            "8 个真实源媒体 SHA256、尺寸和锚点正确" if not image_issues else "；".join(image_issues[:20]),
        )
    )

    task_root = Path(__file__).resolve().parents[1]
    template_path = task_root / "data" / "input_files" / TEMPLATE_NAME
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        template = load_workbook(template_path, data_only=False)
    preserved_issues = preservation_errors(template, workbook)
    results.append(
        criterion(
            "R006",
            not preserved_issues,
            "A1:F9 边框与填充逐格保持模板" if not preserved_issues else "；".join(preserved_issues[:20]),
        )
    )

    rendered_ok, render_evidence = render_check(output_path, expected)
    results.append(criterion("R007", rendered_ok, render_evidence))

    editable_ok = (
        all(isinstance(worksheet[ref].value, str) and worksheet[ref].value for ref in ("C4", "C5", "C6", "B7"))
        and worksheet.protection.sheet is not True
        and len(getattr(worksheet, "_images", [])) == 8
    )
    results.append(
        criterion(
            "R008",
            editable_ok,
            "三层作业和反馈均为可编辑单元格文本" if editable_ok else "内容为空、工作表受保护或被图片替代",
        )
    )

    results.append(criterion("R009", len(content_issues) >= 2, f"源内容错误组数={len(content_issues)}"))
    results.append(
        criterion(
            "R010",
            bool(structure_issues or preserved_issues),
            f"结构错误={len(structure_issues)}，模板外观错误={len(preserved_issues)}",
        )
    )
    results.append(
        criterion(
            "R011",
            bool(layout_issues or image_issues or not rendered_ok),
            f"版式错误={len(layout_issues)}，图片错误={len(image_issues)}，渲染通过={rendered_ok}",
        )
    )
    return build_result(results, output_path)


def build_result(results: list[dict], artifact_path: Path) -> dict:
    present = {item["criterion_id"] for item in results}
    for identifier in CRITERION_SPECS:
        if identifier not in present:
            results.append(criterion(identifier, False, "因交付不存在、损坏或门槛失败，本判据未执行"))
    results.sort(key=lambda item: item["criterion_id"])

    positive_total = sum(item["weight"] for item in results if item["weight"] > 0)
    earned = sum(item["weight"] for item in results if item["weight"] > 0 and item["passed"])
    penalties = sum(item["weight"] for item in results if item["weight"] < 0 and item["passed"])
    reward = max(0.0, min(1.0, (earned + penalties) / positive_total))
    hurdle_failures = [
        item["criterion_id"]
        for item in results
        if item["type"] == "hurdle" and not item["passed"]
    ]
    if hurdle_failures:
        reward = 0.0
    digest = (
        hashlib.sha256(artifact_path.read_bytes()).hexdigest()
        if artifact_path.is_file()
        else hashlib.sha256(b"").hexdigest()
    )
    normalized_results = []
    for item in results:
        score = (0.0 if item["passed"] else 1.0) if item["type"] == "penalty" else item["score"]
        normalized_results.append(
            {
                "criterion_id": item["criterion_id"],
                "score": score,
                "evidence": item["evidence"],
                "method": item["method"],
                "raw": {
                    "description": item["description"],
                    "weight": item["weight"],
                    "type": item["type"],
                    "passed": item["passed"],
                },
            }
        )
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [digest],
        "criterion_results": normalized_results,
        "reward": round(reward, 6),
        "passed": reward >= 0.75,
        "hurdle_failures": hurdle_failures,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path)
    parser.add_argument("--reward-file", type=Path)
    args = parser.parse_args()
    result = evaluate(args.artifact_dir)
    rendered = json.dumps(result, ensure_ascii=False, indent=2)
    print(rendered)
    if args.result:
        args.result.parent.mkdir(parents=True, exist_ok=True)
        args.result.write_text(rendered + "\n", encoding="utf-8")
    if args.reward_file:
        args.reward_file.parent.mkdir(parents=True, exist_ok=True)
        args.reward_file.write_text(f"{result['reward']:.6f}\n", encoding="utf-8")
    raise SystemExit(0 if result["passed"] else 1)


if __name__ == "__main__":
    main()
