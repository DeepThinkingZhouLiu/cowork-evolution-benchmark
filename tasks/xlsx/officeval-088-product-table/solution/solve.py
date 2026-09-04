#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image


EXCEL_TEMPLATE = "选货汇总表_模板2_已填充 (1).xlsx"
OUTPUT = "officeval_088_product_table.xlsx"
PRODUCT_SOURCES = [
    ("云溪", "云溪5.0 D3滚筒洗衣机产品说明文档.docx", (70, 199, 202, 375), 0.75),
    ("统帅", "统帅三桶懒人洗衣机产品说明文档.docx", (232, 158, 471, 500), 0.70),
    ("小天鹅", "小天鹅12KG滚筒洗衣机产品说明文档.docx", (105, 367, 359, 669), 0.84),
    ("米家", "米家8KG滚筒洗衣机产品说明文档.docx", (141, 276, 315, 536), 0.67),
    ("美的", "美的8KG波轮洗衣机产品说明文档.docx", (210, 120, 456, 498), 0.65),
]


def first_media_bytes(path: Path) -> bytes:
    with zipfile.ZipFile(path) as archive:
        names = [name for name in archive.namelist() if name.startswith("word/media/") and not name.endswith("/")]
        if len(names) != 1:
            raise ValueError(f"{path.name}: expected one source image, found {len(names)}")
        return archive.read(names[0])


def parse_product(path: Path) -> dict[str, str | int]:
    text = "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
    def field(label: str) -> str:
        match = re.search(rf"{re.escape(label)}：?\s*([^\n]+)", text)
        if not match:
            raise ValueError(f"missing {label} in {path.name}")
        return re.sub(r"\s+", " ", match.group(1)).strip()

    name = field("产品名称")
    model = field("产品型号")
    price_match = re.search(r"指导价格：[^\n]*?(\d{3,5})元", text)
    if not price_match:
        raise ValueError(f"missing price in {path.name}")
    params_match = re.search(r"四、核心参数(.*?)(?:五、注意事项|$)", text, re.S)
    scene_match = re.search(r"三、适用场景(.*?)(?:四、核心参数|$)", text, re.S)
    params = "；".join(re.sub(r"\s+", "", line).strip("；") for line in (params_match.group(1) if params_match else "").splitlines() if line.strip())
    scene = re.sub(r"\s+", "", scene_match.group(1)).strip("；。 ") if scene_match else ""
    return {"name": name, "model": model, "price": int(price_match.group(1)), "params": f"型号：{model}；{params}", "scene": scene, "source": path.name}


def build(input_dir: Path, output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="officeval-088-xlsx-") as temp:
        temp_root = Path(temp)
        products: list[dict[str, str | int]] = []
        images: dict[str, Path] = {}
        for key, filename, crop, _ratio in PRODUCT_SOURCES:
            source = input_dir / filename
            product = parse_product(source)
            products.append(product)
            raw = temp_root / f"{key}-raw.bin"
            raw.write_bytes(first_media_bytes(source))
            image_path = temp_root / f"{key}.png"
            with Image.open(raw) as image:
                image.convert("RGB").crop(crop).save(image_path, format="PNG")
            images[key] = image_path

        workbook = load_workbook(input_dir / EXCEL_TEMPLATE)
        products_sheet = workbook.active
        products_sheet.title = "Products"
        headers = ["产品名称", "核心参数", "适用场景", "指导价格", "展示图片", "来源文档"]
        for column, value in enumerate(headers, 1):
            products_sheet.cell(1, column).value = value
        widths = {"A": 24, "B": 42, "C": 36, "D": 13, "E": 16, "F": 34}
        for column, width in widths.items():
            products_sheet.column_dimensions[column].width = width
        thin = Side(style="thin", color="000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        header_fill = PatternFill("solid", fgColor="F47B2A")
        for cell in products_sheet[1]:
            cell.fill = header_fill
            cell.font = Font(name="Microsoft YaHei", size=11, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        products_sheet.row_dimensions[1].height = 26
        for row, product in enumerate(products, 2):
            values = [product["name"], product["params"], product["scene"], product["price"], None, product["source"]]
            for column, value in enumerate(values, 1):
                cell = products_sheet.cell(row, column, value)
                cell.font = Font(name="Microsoft YaHei", size=10)
                cell.alignment = Alignment(horizontal="center" if column in (1, 4, 5) else "left", vertical="center", wrap_text=True)
                cell.border = border
                if column == 4:
                    cell.number_format = '0"元"'
            products_sheet.row_dimensions[row].height = 190
            key = next(key for key, filename, _crop, _ratio in PRODUCT_SOURCES if filename == product["source"])
            picture = ExcelImage(str(images[key]))
            picture.height = 190
            picture.width = 190 * next(ratio for _key, _filename, _crop, ratio in PRODUCT_SOURCES if _key == key)
            products_sheet.add_image(picture, f"E{row}")
        products_sheet.freeze_panes = "A2"
        products_sheet.sheet_view.showGridLines = False
        products_sheet.auto_filter.ref = "A1:F6"
        products_sheet.print_title_rows = "1:1"
        products_sheet.page_setup.orientation = "landscape"
        products_sheet.page_setup.fitToWidth = 1
        products_sheet.page_setup.fitToHeight = 0
        products_sheet.sheet_properties.pageSetUpPr.fitToPage = True
        products_sheet.print_area = "A1:F6"

        summary = workbook.create_sheet("Top3 Analysis")
        summary_headers = ["排名", "产品名称", "型号", "指导价格", "来源文档"]
        for column, value in enumerate(summary_headers, 1):
            summary.cell(1, column).value = value
        ranked = sorted(enumerate(products, 2), key=lambda pair: int(pair[1]["price"]), reverse=True)[:3]
        for rank, (source_row, _product) in enumerate(ranked, 1):
            summary.cell(rank + 1, 1).value = rank
            summary.cell(rank + 1, 2).value = f"=Products!A{source_row}"
            summary.cell(rank + 1, 3).value = f'=MID(Products!B{source_row},4,FIND("；",Products!B{source_row})-4)'
            summary.cell(rank + 1, 4).value = f"=Products!D{source_row}"
            summary.cell(rank + 1, 5).value = f"=Products!F{source_row}"
        for column, width in {"A": 9, "B": 28, "C": 32, "D": 13, "E": 42}.items():
            summary.column_dimensions[column].width = width
        for cell in summary[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
            cell.border = border
        for row in summary.iter_rows(min_row=2, max_row=4, min_col=1, max_col=5):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if cell.column == 4:
                    cell.number_format = '0"元"'
        summary.freeze_panes = "A2"
        summary.auto_filter.ref = "A1:E4"
        workbook.calculation.calcMode = "auto"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.save(output_dir / OUTPUT)
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    build(args.input_dir, args.output_dir)


if __name__ == "__main__":
    main()
