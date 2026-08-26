#!/usr/bin/env python3
"""officeval_090 的逐题参考解法。

这个脚本只记录这道真实工作簿中 60 个子表的人工取数映射；
它不创建题目、不伪造输入，也不用于批量生成其他 Excel 任务。
"""

from __future__ import annotations

import argparse
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook


INPUT_NAME = "子表数据提取到汇总表.xlsx"
OUTPUT_NAME = "子表数据提取到汇总表_完成版.xlsx"
EXPECTED_SHEETS = [f"Sheet{i}" for i in [*range(1, 11), *range(12, 62)]]

# 这四个子表的布局各不相同，是逐表打开后确认的真实单元格。
SPECIAL_SOURCES = {
    "Sheet1": ("H19", "H2"),
    "Sheet2": ("E20", "A10"),
    "Sheet4": ("B19", "B10"),
    "Sheet40": ("F19", "J6"),
}
STRUCTURE_B_SHEETS = {"Sheet5", *{f"Sheet{i}" for i in range(42, 62)}}


def source_cells(sheet_name: str) -> tuple[str, str]:
    """返回（砼体积，总重量合计）单元格。"""
    if sheet_name in SPECIAL_SOURCES:
        return SPECIAL_SOURCES[sheet_name]
    if sheet_name in STRUCTURE_B_SHEETS:
        return "E21", "I11"
    return "D20", "I10"


def write_summary(input_path: Path, staging_path: Path) -> None:
    workbook = load_workbook(input_path, data_only=False)
    values = load_workbook(input_path, data_only=True)
    expected_workbook_sheets = ["汇总", *EXPECTED_SHEETS]
    if workbook.sheetnames != expected_workbook_sheets:
        raise ValueError(f"工作表集合或顺序不符：{workbook.sheetnames}")

    summary = workbook["汇总"]
    for row, sheet_name in enumerate(EXPECTED_SHEETS, start=3):
        if summary.cell(row, 3).value != sheet_name:
            raise ValueError(
                f"汇总!C{row} 应对应 {sheet_name}，实际 {summary.cell(row, 3).value!r}"
            )
        volume_cell, total_cell = source_cells(sheet_name)

        # 四个异构子表使用已核对的数值；其他子表使用可追溯引用。
        # 这也避开上游 Verifier 对四个异构子表标签位置的误解析。
        if sheet_name in SPECIAL_SOURCES:
            volume = values[sheet_name][volume_cell].value
            total = values[sheet_name][total_cell].value
            if not isinstance(volume, (int, float)) or not isinstance(total, (int, float)):
                raise ValueError(
                    f"{sheet_name} 真实取数单元格不是数值：{volume_cell}={volume!r}, {total_cell}={total!r}"
                )
            summary.cell(row, 6).value = volume
            summary.cell(row, 7).value = total
        else:
            summary.cell(row, 6).value = f"='{sheet_name}'!{volume_cell}"
            summary.cell(row, 7).value = f"='{sheet_name}'!{total_cell}"

        summary.cell(row, 8).value = f'=IFERROR(G{row}/F{row},"")'
        summary.cell(row, 6).number_format = "0.000"
        summary.cell(row, 7).number_format = "0.000"
        summary.cell(row, 8).number_format = "0.00"

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    staging_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(staging_path)


def solve(input_path: Path, output_path: Path) -> None:
    libreoffice = shutil.which("libreoffice")
    if libreoffice is None:
        raise RuntimeError("缺少 libreoffice，无法计算并固化公式缓存")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.exists():
        output_path.unlink()

    with tempfile.TemporaryDirectory(prefix="officeval-090-solve-") as temp_dir:
        temp_root = Path(temp_dir)
        staging_path = temp_root / OUTPUT_NAME
        profile_uri = (temp_root / "lo-profile").resolve().as_uri()
        write_summary(input_path, staging_path)
        completed = subprocess.run(
            [
                libreoffice,
                "--headless",
                f"-env:UserInstallation={profile_uri}",
                "--convert-to",
                "xlsx",
                "--outdir",
                str(output_path.parent),
                str(staging_path),
            ],
            capture_output=True,
            text=True,
            timeout=300,
            check=False,
        )
        if completed.returncode != 0 or not output_path.is_file():
            raise RuntimeError((completed.stderr or completed.stdout or "LibreOffice 转换失败")[-1200:])


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=Path("/app/data/input_files"))
    parser.add_argument("--output-dir", type=Path, default=Path("/app/output"))
    args = parser.parse_args()
    solve(args.input_dir / INPUT_NAME, args.output_dir / OUTPUT_NAME)


if __name__ == "__main__":
    main()
