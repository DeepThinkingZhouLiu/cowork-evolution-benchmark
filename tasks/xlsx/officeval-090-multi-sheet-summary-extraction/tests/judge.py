#!/usr/bin/env python3
"""officeval_090 的 Harbor Judge。

保留上游逐题 Verifier，同时修正其对四个异构子表标签坐标的
误解析，并补充 61 个工作表完整性、未目标单元格保护和实际渲染检查。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import shutil
import subprocess
import tempfile
import warnings
from pathlib import Path

from openpyxl import load_workbook

import upstream_verifier


OUTPUT_NAME = "子表数据提取到汇总表_完成版.xlsx"
INPUT_NAME = "子表数据提取到汇总表.xlsx"
TASK_ID = "cowork-evo/officeval-090-multi-sheet-summary-extraction"
RUBRIC_VERSION = "omegause-officeval-090-adaptation-v0.1"
JUDGE_VERSION = "officeval-090-adaptation-v0.1"
EXPECTED_CHILD_SHEETS = [f"Sheet{i}" for i in [*range(1, 11), *range(12, 62)]]
EXPECTED_SHEETS = ["汇总", *EXPECTED_CHILD_SHEETS]
SPECIAL_SOURCES = {
    "Sheet1": ("H19", "H2"),
    "Sheet2": ("E20", "A10"),
    "Sheet4": ("B19", "B10"),
    "Sheet40": ("F19", "J6"),
}
STRUCTURE_A_SHEETS = {
    "Sheet3",
    *{f"Sheet{i}" for i in range(6, 11)},
    *{f"Sheet{i}" for i in range(12, 40)},
    "Sheet41",
}
STRUCTURE_B_SHEETS = {"Sheet5", *{f"Sheet{i}" for i in range(42, 62)}}

CRITERION_SPECS = {
    "R001": ("指定 XLSX 可打开，并完整保留 61 个工作表", 2, "hurdle"),
    "R002": ("汇总表第 3–62 行与 60 个子表一一对应", 2, "hurdle"),
    "R003": ("F/G 列 60 行全部取得真实砼体积和总重量合计", 5, "positive"),
    "R004": ("F/G 列使用 0.000 数值格式且不嵌入单位文字", 3, "positive"),
    "R005": ("H3:H62 使用 G/F 公式并使用 0.00 格式", 3, "positive"),
    "R006": ("A 类结构子表正确引用 D20 和 I10", 2, "positive"),
    "R007": ("B 类结构子表正确引用 E21 和 I11", 2, "positive"),
    "R008": ("除汇总 F3:H62 外的原工作簿内容完整保留", 3, "positive"),
    "R009": ("LibreOffice 可实际渲染汇总表首尾数据", 2, "positive"),
    "R010": ("子表内容被修改", -8, "penalty"),
    "R011": ("F/G 列有 10 行及以上取数错误或缺失", -8, "penalty"),
}


def criterion(identifier: str, passed: bool, evidence: str = "") -> dict:
    description, weight, kind = CRITERION_SPECS[identifier]
    return {
        "criterion_id": identifier,
        "description": description,
        "weight": weight,
        "type": kind,
        "passed": bool(passed),
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
    }


def source_cells(sheet_name: str) -> tuple[str, str]:
    if sheet_name in SPECIAL_SOURCES:
        return SPECIAL_SOURCES[sheet_name]
    if sheet_name in STRUCTURE_B_SHEETS:
        return "E21", "I11"
    return "D20", "I10"


def numeric(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(float(value))


def formula_reference(value) -> tuple[str, str] | None:
    if not isinstance(value, str) or not value.startswith("="):
        return None
    match = re.fullmatch(
        r"=\s*'?([^'=!]+)'?!\$?([A-Z]{1,3})\$?(\d+)\s*",
        value,
        flags=re.IGNORECASE,
    )
    if match is None:
        return None
    return match.group(1), f"{match.group(2).upper()}{match.group(3)}"


def values_equal(left, right, tolerance: float = 1e-9) -> bool:
    if numeric(left) and numeric(right):
        return abs(float(left) - float(right)) <= tolerance
    return left == right


def check_source_values(workbook, source_values) -> list[str]:
    errors = []
    summary = workbook["汇总"]
    for row, sheet_name in enumerate(EXPECTED_CHILD_SHEETS, start=3):
        volume_cell, total_cell = source_cells(sheet_name)
        expected_volume = source_values[sheet_name][volume_cell].value
        expected_total = source_values[sheet_name][total_cell].value
        for column, coord, expected, label in (
            (6, volume_cell, expected_volume, "砼体积"),
            (7, total_cell, expected_total, "总重量"),
        ):
            actual = summary.cell(row, column).value
            reference = formula_reference(actual)
            if reference == (sheet_name, coord):
                continue
            if numeric(actual) and numeric(expected) and values_equal(actual, expected):
                continue
            errors.append(
                f"{summary.cell(row, column).coordinate}/{sheet_name} {label}: "
                f"应取 {coord}={expected!r}，实际 {actual!r}"
            )
    return errors


def compare_preserved_cells(source, candidate) -> tuple[list[str], list[str]]:
    errors = []
    child_errors = []
    if source.sheetnames != candidate.sheetnames:
        errors.append(f"工作表列表改变：{source.sheetnames} -> {candidate.sheetnames}")
        child_errors.append("工作表列表改变")
    for sheet_name in source.sheetnames:
        if sheet_name not in candidate.sheetnames:
            continue
        before, after = source[sheet_name], candidate[sheet_name]
        before_merges = {str(item) for item in before.merged_cells.ranges}
        after_merges = {str(item) for item in after.merged_cells.ranges}
        if before_merges != after_merges:
            message = f"{sheet_name} 合并区域改变"
            errors.append(message)
            if sheet_name != "汇总":
                child_errors.append(message)
        max_row = max(before.max_row, after.max_row)
        max_col = max(before.max_column, after.max_column)
        for row in range(1, max_row + 1):
            for column in range(1, max_col + 1):
                if sheet_name == "汇总" and 3 <= row <= 62 and 6 <= column <= 8:
                    continue
                left = before.cell(row, column).value
                right = after.cell(row, column).value
                if values_equal(left, right):
                    continue
                message = f"{sheet_name}!{before.cell(row, column).coordinate}: {left!r} -> {right!r}"
                errors.append(message)
                if sheet_name != "汇总":
                    child_errors.append(message)
    return errors, child_errors


def render_check(path: Path) -> tuple[bool, str]:
    libreoffice = shutil.which("libreoffice")
    pdftotext = shutil.which("pdftotext")
    if libreoffice is None or pdftotext is None:
        return False, "缺少 libreoffice 或 pdftotext"
    with tempfile.TemporaryDirectory(prefix="cowork-evo-090-") as temp_dir:
        temp_root = Path(temp_dir)
        profile_uri = (temp_root / "lo-profile").resolve().as_uri()
        converted = subprocess.run(
            [
                libreoffice,
                "--headless",
                f"-env:UserInstallation={profile_uri}",
                "--convert-to",
                "pdf",
                "--outdir",
                temp_dir,
                str(path),
            ],
            capture_output=True,
            text=True,
            timeout=300,
            check=False,
        )
        pdf_path = temp_root / f"{path.stem}.pdf"
        if converted.returncode != 0 or not pdf_path.is_file():
            return False, (converted.stderr or converted.stdout or "PDF 转换失败")[-1000:]
        extracted = subprocess.run(
            [pdftotext, "-f", "1", "-l", "6", "-layout", str(pdf_path), "-"],
            capture_output=True,
            text=True,
            timeout=120,
            check=False,
        )
        if extracted.returncode != 0:
            return False, (extracted.stderr or "PDF 文字抽取失败")[-1000:]
        normalized = re.sub(r"\s+", "", extracted.stdout)
    required = ["构件材料数据汇总", "PC-GZ-001", "PC-GZ-060", "0.049", "19.006", "56.682", "144.23"]
    missing = [item for item in required if item not in normalized]
    return not missing, f"前 6 个打印页缺失={missing}"


def evaluate(artifact_dir: Path) -> dict:
    output_path = artifact_dir / OUTPUT_NAME
    results = []
    if not output_path.is_file():
        results.append(criterion("R001", False, f"缺少 {OUTPUT_NAME}"))
        return build_result(results, output_path)

    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            workbook = load_workbook(output_path, data_only=False)
            workbook_values = load_workbook(output_path, data_only=True)
    except Exception as exc:
        results.append(criterion("R001", False, f"文件无法打开：{exc}"))
        return build_result(results, output_path)

    sheets_ok = workbook.sheetnames == EXPECTED_SHEETS
    results.append(
        criterion(
            "R001",
            sheets_ok,
            f"工作表数={len(workbook.sheetnames)}，Sheet11={'Sheet11' in workbook.sheetnames}",
        )
    )
    if not sheets_ok:
        return build_result(results, output_path)

    summary = workbook["汇总"]
    actual_mapping = [summary.cell(row, 3).value for row in range(3, 63)]
    mapping_ok = actual_mapping == EXPECTED_CHILD_SHEETS and len(set(actual_mapping)) == 60
    results.append(
        criterion(
            "R002",
            mapping_ok,
            "60 行子表一一对应" if mapping_ok else f"实际映射={actual_mapping}",
        )
    )
    if not mapping_ok:
        return build_result(results, output_path)

    source_path = Path(__file__).resolve().parents[1] / "data" / "input_files" / INPUT_NAME
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        source = load_workbook(source_path, data_only=False)
        source_values = load_workbook(source_path, data_only=True)

    source_errors = check_source_values(workbook, source_values)
    results.append(
        criterion(
            "R003",
            not source_errors,
            "60 行 F/G 均与真实子表一致" if not source_errors else "；".join(source_errors[:20]),
        )
    )

    format_errors = []
    h_errors = []
    for row in range(3, 63):
        for column in (6, 7):
            cell = summary.cell(row, column)
            if cell.number_format != "0.000":
                format_errors.append(f"{cell.coordinate}={cell.number_format!r}")
            if isinstance(cell.value, str) and any(unit in cell.value for unit in ("kg", "m³", "m3")):
                format_errors.append(f"{cell.coordinate}嵌入单位")
        h_cell = summary.cell(row, 8)
        if not upstream_verifier.evaluate_h_formula_equivalent(h_cell.value, row):
            h_errors.append(f"{h_cell.coordinate}={h_cell.value!r}")
        if h_cell.number_format != "0.00":
            h_errors.append(f"{h_cell.coordinate}格式={h_cell.number_format!r}")
    results.append(criterion("R004", not format_errors, "格式正确" if not format_errors else "；".join(format_errors[:20])))
    results.append(criterion("R005", not h_errors, "H3:H62 公式与格式正确" if not h_errors else "；".join(h_errors[:20])))

    upstream_points = upstream_verifier.dimension2_checks(workbook, workbook_values)
    results.append(criterion("R006", upstream_points[3].passed, upstream_points[3].detail or "A 类子表引用正确"))
    results.append(criterion("R007", upstream_points[4].passed, upstream_points[4].detail or "B 类子表引用正确"))

    preservation_errors, child_errors = compare_preserved_cells(source, workbook)
    results.append(
        criterion(
            "R008",
            not preservation_errors,
            "非目标单元格和合并区域完整保留"
            if not preservation_errors
            else "；".join(preservation_errors[:20]),
        )
    )

    rendered, render_evidence = render_check(output_path)
    results.append(criterion("R009", rendered, render_evidence))
    results.append(criterion("R010", bool(child_errors), f"子表修改数={len(child_errors)}"))
    wrong_rows = {error.split("/")[0][1:] for error in source_errors if "/" in error}
    results.append(criterion("R011", len(wrong_rows) >= 10, f"错误行数={len(wrong_rows)}"))
    return build_result(results, output_path)


def build_result(results: list[dict], artifact_path: Path) -> dict:
    present = {item["criterion_id"] for item in results}
    for identifier in CRITERION_SPECS:
        if identifier not in present:
            results.append(criterion(identifier, False, "因交付不存在、损坏或门槛失败，本判据未执行"))
    results.sort(key=lambda item: item["criterion_id"])

    positive_total = sum(item["weight"] for item in results if item["weight"] > 0)
    earned = sum(item["weight"] for item in results if item["weight"] > 0 and item["passed"])
    penalties = sum(item["weight"] for item in results if item["weight"] < 0 and item["passed"])
    reward = max(0.0, min(1.0, (earned + penalties) / positive_total))
    hurdle_failures = [
        item["criterion_id"] for item in results if item["type"] == "hurdle" and not item["passed"]
    ]
    if hurdle_failures:
        reward = 0.0

    digest = hashlib.sha256(artifact_path.read_bytes()).hexdigest() if artifact_path.is_file() else hashlib.sha256(b"").hexdigest()
    normalized = []
    for item in results:
        score = (0.0 if item["passed"] else 1.0) if item["type"] == "penalty" else item["score"]
        normalized.append(
            {
                "criterion_id": item["criterion_id"],
                "score": score,
                "evidence": item["evidence"],
                "method": item["method"],
                "raw": {
                    "description": item["description"],
                    "weight": item["weight"],
                    "type": item["type"],
                    "passed": item["passed"],
                },
            }
        )
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [digest],
        "criterion_results": normalized,
        "reward": round(reward, 6),
        "passed": reward >= 0.75,
        "hurdle_failures": hurdle_failures,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path)
    parser.add_argument("--reward-file", type=Path)
    args = parser.parse_args()
    result = evaluate(args.artifact_dir)
    rendered = json.dumps(result, ensure_ascii=False, indent=2)
    print(rendered)
    if args.result:
        args.result.parent.mkdir(parents=True, exist_ok=True)
        args.result.write_text(rendered + "\n", encoding="utf-8")
    if args.reward_file:
        args.reward_file.parent.mkdir(parents=True, exist_ok=True)
        args.reward_file.write_text(f"{result['reward']:.6f}\n", encoding="utf-8")
    raise SystemExit(0 if result["passed"] else 1)


if __name__ == "__main__":
    main()
