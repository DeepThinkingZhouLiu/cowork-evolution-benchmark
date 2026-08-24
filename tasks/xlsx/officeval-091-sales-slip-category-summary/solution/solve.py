#!/usr/bin/env python3
"""officeval_091 的逐题参考解法。

销售单中的客户、商品和 1600kg 重量由人工查看本题 image_001.png 后转录；
目标行和 57 个客户区域由人工检查本题固定工作簿后确定。代码只用于记录这道
真实任务的可回放修改，不是通用 OCR、透视表生成器或批量制题脚本。
"""

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


INPUT_NAME = "2026年区域渠道销量统计表.xlsx"
IMAGE_NAME = "image_001.png"
OUTPUT_NAME = "2026年区域渠道销量统计表_整理完成.xlsx"
INPUT_SHA256 = "afcc3774aebb6962b279ee2fa3461fe4a8fbdbe7759cca5a931ce7edbb5fa5be"
IMAGE_SHA256 = "ece31a54f2498779bb24fb2144e2478e3fed76a265b6d06bc81851ae2720420b"

# 人工读取真实销售单：连塘合作养殖社，育成期配合饲料，1600kg。
SCREENSHOT_CUSTOMER = "连塘合作养殖社"
SCREENSHOT_PRODUCT = "育成期配合饲料"
SCREENSHOT_WEIGHT_KG = 1600
SCREENSHOT_WEIGHT_TONS = 1.6

CATEGORY_HEADERS = ["开口料", "保育/育肥料", "蛋禽料", "水禽料", "沉水鱼料", "浮水鱼料"]

# 每段均为逐块检查“12月”工作表后确定的客户区：
# (客户列, 第一类列, 起始行, 结束行)，不包含表头和小计/合计行。
CUSTOMER_SEGMENTS = [
    (2, 4, 5, 10),
    (11, 13, 5, 7),
    (2, 4, 19, 24),
    (11, 13, 19, 53),
    (2, 4, 40, 46),
]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def validate_fixed_inputs(workbook_path: Path, image_path: Path) -> None:
    if sha256(workbook_path) != INPUT_SHA256:
        raise ValueError("工作簿不是本题固定的真实输入版本")
    if sha256(image_path) != IMAGE_SHA256:
        raise ValueError("销售单图片不是本题固定的真实输入版本")


def customer_sources(worksheet) -> list[tuple[str, int, int]]:
    """返回（客户名、来源行、六分类起始列），严格按原表视觉顺序。"""

    result = []
    for customer_column, category_start, start_row, end_row in CUSTOMER_SEGMENTS:
        for row in range(start_row, end_row + 1):
            name = worksheet.cell(row, customer_column).value
            if not isinstance(name, str) or not name.strip():
                raise ValueError(f"客户源单元格为空：{worksheet.cell(row, customer_column).coordinate}")
            result.append((name.strip(), row, category_start))
    if len(result) != 57 or len({item[0] for item in result}) != 57:
        raise ValueError("本题应有 57 个互不重复的真实客户")
    return result


def apply_screenshot_record(worksheet) -> None:
    if worksheet["K41"].value != SCREENSHOT_CUSTOMER:
        raise ValueError(f"销售单客户未定位到 12月!K41：{worksheet['K41'].value!r}")
    if abs(float(worksheet["N41"].value) - 0.92) > 1e-9:
        raise ValueError(f"12月!N41 原值应为 0.92 吨：{worksheet['N41'].value!r}")
    if worksheet["L41"].value != "=SUM(M41:R41)":
        raise ValueError("12月!L41 原汇总公式与本题预期不符")
    if any(worksheet[ref].value not in (None, "") for ref in ("O41", "P41", "Q41", "R41")):
        raise ValueError("12月!O41:R41 原本应为空")

    # “育成期配合饲料”归入保育/育肥料；1600kg ÷ 1000 = 1.60t。
    worksheet["N41"] = round(0.92 + SCREENSHOT_WEIGHT_TONS, 2)
    worksheet["N41"].number_format = "0.00"


def style_summary(summary) -> None:
    dark_blue = "1F4E78"
    medium_blue = "5B9BD5"
    light_blue = "DDEBF7"
    light_green = "E2F0D9"
    grid = Side(style="thin", color="D9E2F3")
    border = Border(left=grid, right=grid, top=grid, bottom=grid)

    summary.merge_cells("A1:H1")
    summary["A1"] = "2026年区域渠道销量客户分类汇总"
    summary["A1"].font = Font(name="微软雅黑", size=16, bold=True, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor=dark_blue)
    summary["A1"].alignment = Alignment(horizontal="center", vertical="center")
    summary.row_dimensions[1].height = 30

    summary.merge_cells("A2:H2")
    summary["A2"] = (
        f"来源：12月明细；销售单 {SCREENSHOT_CUSTOMER} / {SCREENSHOT_PRODUCT} / "
        f"{SCREENSHOT_WEIGHT_KG}kg 已换算为 {SCREENSHOT_WEIGHT_TONS:.2f} 吨后计入保育/育肥料。"
    )
    summary["A2"].font = Font(name="微软雅黑", size=10, color="404040")
    summary["A2"].fill = PatternFill("solid", fgColor=light_blue)
    summary["A2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    summary.row_dimensions[2].height = 32

    headers = ["客户名称", "数量(吨)", *CATEGORY_HEADERS]
    for column, value in enumerate(headers, start=1):
        cell = summary.cell(4, column, value)
        cell.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=medium_blue)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    summary.row_dimensions[4].height = 28

    for row in range(5, 62):
        if row % 2 == 0:
            for column in range(1, 9):
                summary.cell(row, column).fill = PatternFill("solid", fgColor="F7FBFF")
        for column in range(1, 9):
            cell = summary.cell(row, column)
            cell.font = Font(name="微软雅黑", size=9.5, color="000000")
            cell.alignment = Alignment(
                horizontal="left" if column == 1 else "right",
                vertical="center",
            )
            cell.border = border
        summary.row_dimensions[row].height = 21

    total_row = 62
    for column in range(1, 9):
        cell = summary.cell(total_row, column)
        cell.font = Font(name="微软雅黑", size=10, bold=True, color="000000")
        cell.fill = PatternFill("solid", fgColor=light_green)
        cell.border = border
        cell.alignment = Alignment(
            horizontal="left" if column == 1 else "right",
            vertical="center",
        )
    summary.row_dimensions[total_row].height = 23

    widths = {"A": 24, "B": 13, "C": 12, "D": 16, "E": 12, "F": 12, "G": 12, "H": 12}
    for column, width in widths.items():
        summary.column_dimensions[column].width = width
    for column in range(2, 9):
        for row in range(5, 63):
            summary.cell(row, column).number_format = "0.00"

    summary.freeze_panes = "A5"
    summary.auto_filter.ref = "A4:H61"
    summary.sheet_view.showGridLines = False
    summary.print_title_rows = "1:4"
    summary.print_area = "A1:H62"
    summary.page_setup.orientation = "landscape"
    summary.page_setup.fitToWidth = 1
    summary.page_setup.fitToHeight = 0
    summary.sheet_properties.pageSetUpPr.fitToPage = True
    summary.page_margins.left = 0.25
    summary.page_margins.right = 0.25
    summary.page_margins.top = 0.45
    summary.page_margins.bottom = 0.45


def build_summary(workbook, source) -> None:
    if "客户分类汇总" in workbook.sheetnames:
        del workbook["客户分类汇总"]
    summary = workbook.create_sheet("客户分类汇总")
    sources = customer_sources(source)

    for output_row, (customer, source_row, category_start) in enumerate(sources, start=5):
        summary.cell(output_row, 1, customer)
        summary.cell(output_row, 2, f"=SUM(C{output_row}:H{output_row})")
        for category_offset in range(6):
            source_coordinate = f"{get_column_letter(category_start + category_offset)}{source_row}"
            summary.cell(
                output_row,
                3 + category_offset,
                f"='12月'!{source_coordinate}",
            )

    summary["A62"] = "合计"
    for column in range(2, 9):
        letter = get_column_letter(column)
        summary.cell(62, column, f"=SUM({letter}5:{letter}61)")
    style_summary(summary)


def solve(input_dir: Path, output_dir: Path) -> None:
    workbook_path = input_dir / INPUT_NAME
    image_path = input_dir / IMAGE_NAME
    validate_fixed_inputs(workbook_path, image_path)
    workbook = load_workbook(workbook_path, data_only=False)
    if workbook.sheetnames != ["12月"]:
        raise ValueError(f"原工作簿工作表不符：{workbook.sheetnames!r}")
    source = workbook["12月"]
    apply_screenshot_record(source)
    build_summary(workbook, source)

    if workbook.sheetnames != ["12月", "客户分类汇总"]:
        raise ValueError("输出必须且只能包含两张工作表")
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"

    output_dir.mkdir(parents=True, exist_ok=True)
    workbook.save(output_dir / OUTPUT_NAME)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=Path("/app/data/input_files"))
    parser.add_argument("--output-dir", type=Path, default=Path("/app/output"))
    args = parser.parse_args()
    solve(args.input_dir, args.output_dir)


if __name__ == "__main__":
    main()
