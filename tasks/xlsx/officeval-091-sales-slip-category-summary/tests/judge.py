#!/usr/bin/env python3
"""officeval_091 的 Harbor Judge。

逐项核对销售单 1600kg -> 1.60t 的落表位置、57 个真实客户与六类来源公式、
原“12月”表的保留性，以及 LibreOffice 重算后的 741.56 吨汇总。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import subprocess
import tempfile
import warnings
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


OUTPUT_NAME = "2026年区域渠道销量统计表_整理完成.xlsx"
INPUT_NAME = "2026年区域渠道销量统计表.xlsx"
TASK_ID = "cowork-evo/officeval-091-sales-slip-category-summary"
RUBRIC_VERSION = "omegause-officeval-091-adaptation-v0.1"
JUDGE_VERSION = "officeval-091-adaptation-v0.1"
EXPECTED_SHEETS = ["12月", "客户分类汇总"]
HEADERS = ["客户名称", "数量(吨)", "开口料", "保育/育肥料", "蛋禽料", "水禽料", "沉水鱼料", "浮水鱼料"]

CRITERION_SPECS = {
    "R001": ("指定 XLSX 可打开，且仅含 12月和客户分类汇总两张工作表", 2, "hurdle"),
    "R002": ("销售单 1600kg 正确换算为 1.60t 并计入连塘合作养殖社", 6, "positive"),
    "R003": ("汇总表包含 57 个真实且互不重复的客户和六类字段", 5, "positive"),
    "R004": ("57 行六分类和总量均以正确公式引用 12月真实明细", 7, "positive"),
    "R005": ("除 N41 外，原 12月表的值、公式、样式和结构完整保留", 4, "positive"),
    "R006": ("汇总表标题、表头、列宽、冻结、筛选、格式和分页设置完整", 4, "positive"),
    "R007": ("LibreOffice 重算和渲染后 2.52 与 741.56 等关键结果可读", 3, "positive"),
    "R008": ("汇总公式可编辑且工作簿启用打开时自动重算", 2, "positive"),
    "R009": ("销售单未录入、kg/吨换算错误或商品归类错误", -8, "penalty"),
    "R010": ("任一客户、分类或来源公式错配", -8, "penalty"),
    "R011": ("原 12月表除目标单元格外被修改", -8, "penalty"),
    "R012": ("汇总表关键版式或实际重算渲染失败", -6, "penalty"),
}


def criterion(identifier: str, passed: bool, evidence: str = "") -> dict:
    description, weight, kind = CRITERION_SPECS[identifier]
    return {
        "criterion_id": identifier,
        "description": description,
        "weight": weight,
        "type": kind,
        "passed": bool(passed),
        "score": 1.0 if passed else 0.0,
        "evidence": evidence,
        "method": "deterministic",
    }


def normalized_formula(value) -> str:
    if not isinstance(value, str):
        return ""
    return value.replace("$", "").replace("'", "").replace(" ", "").upper()


def numeric_equal(left, right, tolerance: float = 1e-9) -> bool:
    if isinstance(left, bool) or isinstance(right, bool):
        return left == right
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return abs(float(left) - float(right)) <= tolerance
    return left == right


def expected_sources(source, expected: dict) -> list[tuple[str, int, int]]:
    rows = []
    for segment in expected["customer_segments"]:
        for row in range(segment["start_row"], segment["end_row"] + 1):
            customer = source.cell(row, segment["customer_column"]).value
            rows.append((customer, row, segment["category_start_column"]))
    return rows


def screenshot_errors(workbook, expected: dict) -> list[str]:
    source = workbook["12月"]
    slip = expected["screenshot"]
    errors = []
    if source["K41"].value != slip["customer"]:
        errors.append(f"K41 客户={source['K41'].value!r}")
    if not numeric_equal(source[slip["target_cell"]].value, slip["final_tons"]):
        errors.append(f"{slip['target_cell']}={source[slip['target_cell']].value!r}，期望={slip['final_tons']}")
    if normalized_formula(source["L41"].value) != "=SUM(M41:R41)":
        errors.append(f"L41 公式={source['L41'].value!r}")
    for ref in ("O41", "P41", "Q41", "R41"):
        if source[ref].value not in (None, ""):
            errors.append(f"{ref} 不应有值：{source[ref].value!r}")
    if source["K57"].value != "合计":
        errors.append(f"K57={source['K57'].value!r}")
    return errors


def customer_errors(source, summary, expected: dict) -> list[str]:
    errors = []
    actual_headers = [summary.cell(4, column).value for column in range(1, 9)]
    if actual_headers != HEADERS:
        errors.append(f"表头={actual_headers!r}")
    sources = expected_sources(source, expected)
    if len(sources) != expected["customer_count"]:
        errors.append(f"隐藏来源客户数={len(sources)}")
    actual_names = [summary.cell(row, 1).value for row in range(5, 62)]
    expected_names = [item[0] for item in sources]
    if actual_names != expected_names:
        for index, (actual, wanted) in enumerate(zip(actual_names, expected_names), start=5):
            if actual != wanted:
                errors.append(f"A{index}={actual!r}，期望={wanted!r}")
    if len(set(actual_names)) != 57:
        errors.append(f"汇总客户去重数={len(set(actual_names))}，期望=57")
    if summary.max_row != 62 or summary.max_column < 8:
        errors.append(f"汇总尺寸={summary.max_row}×{summary.max_column}")
    return errors


def formula_errors(source, summary, expected: dict) -> list[str]:
    errors = []
    sources = expected_sources(source, expected)
    for output_row, (_customer, source_row, category_start) in enumerate(sources, start=5):
        if normalized_formula(summary.cell(output_row, 2).value) != f"=SUM(C{output_row}:H{output_row})":
            errors.append(f"B{output_row} 总量公式={summary.cell(output_row, 2).value!r}")
        for offset in range(6):
            source_ref = f"{get_column_letter(category_start + offset)}{source_row}"
            wanted = f"=12月!{source_ref}".upper()
            actual = normalized_formula(summary.cell(output_row, 3 + offset).value)
            if actual != wanted:
                errors.append(
                    f"{summary.cell(output_row, 3 + offset).coordinate}={summary.cell(output_row, 3 + offset).value!r}，"
                    f"期望引用 12月!{source_ref}"
                )
    if summary["A62"].value != "合计":
        errors.append(f"A62={summary['A62'].value!r}")
    for column in range(2, 9):
        letter = get_column_letter(column)
        wanted = f"=SUM({letter}5:{letter}61)"
        if normalized_formula(summary.cell(62, column).value) != wanted:
            errors.append(f"{letter}62={summary.cell(62, column).value!r}")
    return errors


def preservation_errors(original, candidate) -> list[str]:
    errors = []
    if {str(item) for item in original.merged_cells.ranges} != {str(item) for item in candidate.merged_cells.ranges}:
        errors.append("12月 合并区域改变")
    if (original.max_row, original.max_column) != (candidate.max_row, candidate.max_column):
        errors.append(
            f"12月 尺寸改变：{original.max_row}×{original.max_column} -> "
            f"{candidate.max_row}×{candidate.max_column}"
        )
    for row in range(1, max(original.max_row, candidate.max_row) + 1):
        for column in range(1, max(original.max_column, candidate.max_column) + 1):
            before = original.cell(row, column)
            after = candidate.cell(row, column)
            if before.coordinate == "N41":
                if before.style_id != after.style_id:
                    errors.append("N41 原单元格样式改变")
                continue
            if not numeric_equal(before.value, after.value):
                errors.append(f"{before.coordinate}: {before.value!r} -> {after.value!r}")
            if before.style_id != after.style_id:
                errors.append(f"{before.coordinate} 样式改变：{before.style_id} -> {after.style_id}")
    original_rows = {
        index: (item.height, item.hidden, item.outlineLevel)
        for index, item in original.row_dimensions.items()
    }
    candidate_rows = {
        index: (item.height, item.hidden, item.outlineLevel)
        for index, item in candidate.row_dimensions.items()
    }
    if original_rows != candidate_rows:
        errors.append("12月 行高或隐藏设置改变")
    original_columns = {
        key: (item.width, item.hidden, item.outlineLevel)
        for key, item in original.column_dimensions.items()
    }
    candidate_columns = {
        key: (item.width, item.hidden, item.outlineLevel)
        for key, item in candidate.column_dimensions.items()
    }
    if original_columns != candidate_columns:
        errors.append("12月 列宽或隐藏设置改变")
    return errors


def presentation_errors(summary) -> list[str]:
    errors = []
    merges = {str(item) for item in summary.merged_cells.ranges}
    if merges != {"A1:H1", "A2:H2"}:
        errors.append(f"汇总合并区域={sorted(merges)}")
    if summary["A1"].value != "2026年区域渠道销量客户分类汇总":
        errors.append("汇总标题错误")
    if summary.freeze_panes != "A5":
        errors.append(f"冻结窗格={summary.freeze_panes!r}")
    if summary.auto_filter.ref != "A4:H61":
        errors.append(f"筛选范围={summary.auto_filter.ref!r}")
    widths = {"A": 24, "B": 13, "C": 12, "D": 16, "E": 12, "F": 12, "G": 12, "H": 12}
    for column, wanted in widths.items():
        actual = summary.column_dimensions[column].width
        if actual is None or abs(float(actual) - wanted) > 0.05:
            errors.append(f"{column}列宽={actual}，期望={wanted}")
    if summary.page_setup.orientation != "landscape" or summary.page_setup.fitToWidth != 1:
        errors.append("打印设置未保持横向单页宽")
    for row in range(5, 63):
        for column in range(2, 9):
            if summary.cell(row, column).number_format != "0.00":
                errors.append(f"{summary.cell(row, column).coordinate} 数值格式错误")
                if len(errors) >= 20:
                    return errors
    return errors


def render_check(path: Path) -> tuple[bool, str]:
    libreoffice = shutil.which("libreoffice")
    pdftotext = shutil.which("pdftotext")
    pdfinfo = shutil.which("pdfinfo")
    if libreoffice is None or pdftotext is None or pdfinfo is None:
        return False, "缺少 libreoffice、pdftotext 或 pdfinfo"
    with tempfile.TemporaryDirectory(prefix="cowork-evo-091-") as temp_dir:
        temp_root = Path(temp_dir)
        profile_uri = (temp_root / "lo-profile").resolve().as_uri()
        converted = subprocess.run(
            [
                libreoffice,
                "--headless",
                f"-env:UserInstallation={profile_uri}",
                "--convert-to",
                "pdf",
                "--outdir",
                temp_dir,
                str(path),
            ],
            capture_output=True,
            text=True,
            timeout=300,
            check=False,
        )
        pdf_path = temp_root / f"{path.stem}.pdf"
        if converted.returncode != 0 or not pdf_path.is_file():
            return False, (converted.stderr or converted.stdout or "PDF 转换失败")[-1000:]
        info = subprocess.run([pdfinfo, str(pdf_path)], capture_output=True, text=True, timeout=120, check=False)
        extracted = subprocess.run(
            [pdftotext, "-raw", str(pdf_path), "-"],
            capture_output=True,
            text=True,
            timeout=120,
            check=False,
        )
        if info.returncode != 0 or extracted.returncode != 0:
            return False, (info.stderr or extracted.stderr or "PDF 检查失败")[-1000:]
        match = re.search(r"^Pages:\s+(\d+)", info.stdout, flags=re.MULTILINE)
        pages = int(match.group(1)) if match else 0
        rendered = re.sub(r"\s+", "", extracted.stdout)
    required = ["客户分类汇总", "连塘合作养殖社", "2.52", "741.56", "保育/育肥料", "浮水鱼料"]
    missing = [item for item in required if re.sub(r"\s+", "", item) not in rendered]
    return 7 <= pages <= 10 and not missing, f"页数={pages}，渲染缺失={missing}"


def evaluate(artifact_dir: Path) -> dict:
    output_path = artifact_dir / OUTPUT_NAME
    results = []
    if not output_path.is_file():
        results.append(criterion("R001", False, f"缺少 {OUTPUT_NAME}"))
        return build_result(results, output_path)
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            workbook = load_workbook(output_path, data_only=False)
    except Exception as exc:
        results.append(criterion("R001", False, f"文件无法打开：{exc}"))
        return build_result(results, output_path)

    sheets_ok = workbook.sheetnames == EXPECTED_SHEETS
    results.append(criterion("R001", sheets_ok, f"工作表={workbook.sheetnames}"))
    if not sheets_ok:
        return build_result(results, output_path)

    task_root = Path(__file__).resolve().parents[1]
    expected = json.loads(Path(__file__).with_name("expected-sources.json").read_text(encoding="utf-8"))
    source = workbook["12月"]
    summary = workbook["客户分类汇总"]

    slip_issues = screenshot_errors(workbook, expected)
    results.append(
        criterion("R002", not slip_issues, "1600kg -> 1.60t，N41=2.52" if not slip_issues else "；".join(slip_issues))
    )

    customer_issues = customer_errors(source, summary, expected)
    results.append(
        criterion("R003", not customer_issues, "57 个客户顺序和六类字段正确" if not customer_issues else "；".join(customer_issues[:20]))
    )

    formulas_issues = formula_errors(source, summary, expected)
    results.append(
        criterion("R004", not formulas_issues, "57×7 个来源/总量公式及合计公式正确" if not formulas_issues else "；".join(formulas_issues[:20]))
    )

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        original_workbook = load_workbook(task_root / "data" / "input_files" / INPUT_NAME, data_only=False)
    preserved_issues = preservation_errors(original_workbook["12月"], source)
    results.append(
        criterion("R005", not preserved_issues, "除 N41 外原表逐格保留" if not preserved_issues else "；".join(preserved_issues[:20]))
    )

    visual_issues = presentation_errors(summary)
    results.append(
        criterion("R006", not visual_issues, "汇总标题、表头、格式、筛选和分页完整" if not visual_issues else "；".join(visual_issues[:20]))
    )

    rendered_ok, render_evidence = render_check(output_path)
    results.append(criterion("R007", rendered_ok, render_evidence))

    calc = workbook.calculation
    editable_ok = (
        calc.calcMode == "auto"
        and bool(calc.fullCalcOnLoad)
        and bool(calc.forceFullCalc)
        and all(isinstance(summary.cell(row, column).value, str) and summary.cell(row, column).value.startswith("=") for row in range(5, 62) for column in range(2, 9))
        and summary.protection.sheet is not True
    )
    results.append(
        criterion("R008", editable_ok, "公式可编辑并启用打开时自动重算" if editable_ok else "公式、保护或重算设置不符合")
    )

    results.append(criterion("R009", bool(slip_issues), f"销售单落表错误={len(slip_issues)}"))
    results.append(
        criterion("R010", bool(customer_issues or formulas_issues), f"客户错误={len(customer_issues)}，公式错误={len(formulas_issues)}")
    )
    results.append(criterion("R011", bool(preserved_issues), f"原表保护错误={len(preserved_issues)}"))
    results.append(
        criterion("R012", bool(visual_issues or not rendered_ok), f"版式错误={len(visual_issues)}，渲染通过={rendered_ok}")
    )
    return build_result(results, output_path)


def build_result(results: list[dict], artifact_path: Path) -> dict:
    present = {item["criterion_id"] for item in results}
    for identifier in CRITERION_SPECS:
        if identifier not in present:
            results.append(criterion(identifier, False, "因交付不存在、损坏或门槛失败，本判据未执行"))
    results.sort(key=lambda item: item["criterion_id"])
    positive_total = sum(item["weight"] for item in results if item["weight"] > 0)
    earned = sum(item["weight"] for item in results if item["weight"] > 0 and item["passed"])
    penalties = sum(item["weight"] for item in results if item["weight"] < 0 and item["passed"])
    reward = max(0.0, min(1.0, (earned + penalties) / positive_total))
    hurdle_failures = [item["criterion_id"] for item in results if item["type"] == "hurdle" and not item["passed"]]
    if hurdle_failures:
        reward = 0.0
    digest = hashlib.sha256(artifact_path.read_bytes()).hexdigest() if artifact_path.is_file() else hashlib.sha256(b"").hexdigest()
    normalized_results = []
    for item in results:
        score = (0.0 if item["passed"] else 1.0) if item["type"] == "penalty" else item["score"]
        normalized_results.append(
            {
                "criterion_id": item["criterion_id"],
                "score": score,
                "evidence": item["evidence"],
                "method": item["method"],
                "raw": {
                    "description": item["description"],
                    "weight": item["weight"],
                    "type": item["type"],
                    "passed": item["passed"],
                },
            }
        )
    return {
        "task_id": TASK_ID,
        "rubric_version": RUBRIC_VERSION,
        "judge_version": JUDGE_VERSION,
        "artifact_sha256": [digest],
        "criterion_results": normalized_results,
        "reward": round(reward, 6),
        "passed": reward >= 0.75,
        "hurdle_failures": hurdle_failures,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path)
    args = parser.parse_args()
    result = evaluate(args.artifact_dir)
    rendered = json.dumps(result, ensure_ascii=False, indent=2)
    print(rendered)
    if args.result:
        args.result.parent.mkdir(parents=True, exist_ok=True)
        args.result.write_text(rendered + "\n", encoding="utf-8")
    raise SystemExit(0 if result["passed"] else 1)


if __name__ == "__main__":
    main()
