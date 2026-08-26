#!/usr/bin/env python3
"""officeval_092 的逐题参考解法。

本脚本只处理这一道真实来源任务。表格内容来自输入工作簿内嵌截图，
行号、公式、样式与分页要求来自上游 Query 和 Rubric，不作为通用制题器使用。
"""

from __future__ import annotations

import argparse
import math
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break, RowBreak


INPUT_NAME = "智慧农业示范园建设项目报价表.xlsx"
OUTPUT_NAME = "智慧农业示范园建设项目报价表_完成版.xlsx"
MAIN_SHEET = "智慧农业报价表"

HEADERS = [
    "序号",
    "项目名称",
    "项目特征描述",
    "工程量",
    "单位",
    "综合单价（元）",
    "合价（元）",
    "施工/供应说明",
    "使用部位",
    "备注",
]

ENVIRONMENT_RECORDS = [
    [1, "环流风机", "低噪声轴流，防潮电机", 36.00, "台", 1850.00, "设备供应", "温室内部", ""],
    [2, "湿帘纸芯", "150mm厚高效蒸发湿帘", 220.00, "㎡", 165.00, "设备供应", "湿帘墙", ""],
    [3, "LED补光灯", "植物生长光谱，防水等级IP65", 168.00, "盏", 680.00, "设备供应", "育苗区", ""],
    [4, "遮阳幕布", "外白内黑节能幕，阻燃处理", 8600.00, "㎡", 22.00, "设备供应", "顶部遮阳", ""],
    [5, "电动卷膜器", "24V直流，带限位保护", 44.00, "套", 960.00, "设备供应", "侧窗通风", ""],
]

IOT_RECORDS = [
    [1, "温湿度传感器", "温度、湿度双测点，防水外壳", 28.00, "支", 520.00, "设备供应", "栽培区", ""],
    [2, "CO₂传感器", "NDIR原理，量程0-5000ppm", 12.00, "支", 1680.00, "设备供应", "生产区", ""],
    [3, "土壤/基质水分传感器", "电容式，带温度补偿", 24.00, "支", 780.00, "设备供应", "种植槽", ""],
    [4, "LoRa采集网关", "工业级，支持断点续传", 6.00, "台", 3900.00, "设备供应", "控制间", ""],
    [5, "生产大屏看板", "55英寸显示屏，含支架与信号盒", 2.00, "套", 6800.00, "设备供应", "运营中心", ""],
]

TABLE_BLOCKS = [
    (7, 20),
    (27, 31),
    (34, 38),
    (41, 46),
    (49, 54),
    (61, 66),
    (70, 76),
    (80, 86),
    (90, 96),
    (100, 106),
    (110, 116),
    (120, 124),
]

PAGE_BREAK_ROWS = [23, 32, 39, 47, 58, 68, 78, 88, 98, 108, 118]


def copy_cell_style(source, target) -> None:
    """复制单元格显示样式，但不复制内容。"""

    target._style = copy(source._style)
    target.number_format = source.number_format
    target.protection = copy(source.protection)


def copy_row_style(ws, source_row: int, target_row: int) -> None:
    for column in range(1, 11):
        copy_cell_style(ws.cell(source_row, column), ws.cell(target_row, column))
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def write_equipment_table(
    ws,
    *,
    title_row: int,
    title: str,
    subtotal_title: str,
    records: list[list[object]],
) -> None:
    """把截图中的一张清单写成可编辑表格。"""

    header_row = title_row + 1
    first_data_row = title_row + 2
    subtotal_row = first_data_row + len(records)

    # 从相邻的真实表格复制样式，避免重新设计整份工作簿。
    copy_row_style(ws, 59, title_row)
    copy_row_style(ws, 60, header_row)
    for offset in range(len(records)):
        copy_row_style(ws, 61 if offset % 2 == 0 else 62, first_data_row + offset)
    copy_row_style(ws, 66, subtotal_row)

    ws.merge_cells(start_row=title_row, start_column=1, end_row=title_row, end_column=10)
    ws.cell(title_row, 1).value = title

    for column, value in enumerate(HEADERS, 1):
        ws.cell(header_row, column).value = value

    for offset, record in enumerate(records):
        row = first_data_row + offset
        number, name, description, quantity, unit, unit_price, supply, location, remark = record
        values = [
            number,
            name,
            description,
            quantity,
            unit,
            unit_price,
            f"=D{row}*F{row}",
            supply,
            location,
            remark,
        ]
        for column, value in enumerate(values, 1):
            ws.cell(row, column).value = value
        for column in (4, 6, 7):
            ws.cell(row, column).number_format = "#,##0.00"

    ws.merge_cells(start_row=subtotal_row, start_column=1, end_row=subtotal_row, end_column=6)
    ws.cell(subtotal_row, 1).value = subtotal_title
    ws.cell(subtotal_row, 7).value = f"=SUM(G{first_data_row}:G{subtotal_row - 1})"
    ws.cell(subtotal_row, 7).number_format = "#,##0.00"


def apply_uniform_table_typography(ws) -> None:
    """按原 Query 统一表内字体、对齐和换行。"""

    for row_start, row_end in TABLE_BLOCKS:
        for row in range(row_start, row_end + 1):
            for column in range(1, 11):
                cell = ws.cell(row, column)
                if isinstance(cell, MergedCell) or cell.value is None:
                    continue
                cell.font = copy(cell.font)
                cell.font = Font(
                    name="微软雅黑",
                    size=10,
                    bold=cell.font.bold,
                    italic=cell.font.italic,
                    color="000000",
                    underline=cell.font.underline,
                )
                cell.alignment = copy(cell.alignment)
                cell.alignment = Alignment(
                    horizontal=cell.alignment.horizontal,
                    vertical="center",
                    text_rotation=cell.alignment.text_rotation,
                    wrap_text=True,
                    shrink_to_fit=cell.alignment.shrink_to_fit,
                    indent=cell.alignment.indent,
                )


def _visual_width(text: str) -> float:
    return sum(1.4 if ord(character) > 127 else 0.72 for character in text)


def cell_available_width(ws, row: int, column: int) -> float:
    """返回单元格或其合并区域的可用列宽。"""

    for merged_range in ws.merged_cells.ranges:
        if (
            merged_range.min_row <= row <= merged_range.max_row
            and merged_range.min_col <= column <= merged_range.max_col
        ):
            return sum(
                float(ws.column_dimensions[get_column_letter(col)].width or 8.43)
                for col in range(merged_range.min_col, merged_range.max_col + 1)
            )
    return float(ws.column_dimensions[get_column_letter(column)].width or 8.43)


def required_row_height(ws, row: int) -> float:
    required_lines = 1
    for column in range(1, 11):
        cell = ws.cell(row, column)
        if isinstance(cell, MergedCell) or cell.value is None:
            continue
        if isinstance(cell.value, str) and cell.value.startswith("="):
            continue
        if not cell.alignment.wrap_text:
            continue
        available = max(cell_available_width(ws, row, column) - 1.6, 1.0)
        lines = 0
        for logical_line in str(cell.value).splitlines() or [""]:
            lines += max(1, math.ceil(_visual_width(logical_line) / available))
        required_lines = max(required_lines, lines)
    return max(18.0, required_lines * 13.5 + 4.0)


def adjust_row_heights(ws) -> None:
    """依据实际文本、合并区域和列宽逐行计算最低行高。"""

    for row in range(1, 125):
        has_content = any(
            not isinstance(ws.cell(row, column), MergedCell)
            and ws.cell(row, column).value is not None
            for column in range(1, 11)
        )
        if has_content:
            ws.row_dimensions[row].height = max(
                float(ws.row_dimensions[row].height or 15.0),
                required_row_height(ws, row),
            )

    # 表名、分组标题和表头分别统一最低高度。
    for row in (1, 5, 24, 59, 69, 79, 89, 99, 109, 119):
        ws.row_dimensions[row].height = max(float(ws.row_dimensions[row].height or 0), 28.0)
    for row in (26, 33, 40, 48):
        ws.row_dimensions[row].height = max(float(ws.row_dimensions[row].height or 0), 26.0)
    for row in (6, 25, 60, 70, 80, 90, 100, 110, 120):
        ws.row_dimensions[row].height = max(float(ws.row_dimensions[row].height or 0), 24.0)

    # 六张结构相同的物资清单使用一致的数据行和小计行高度。
    repeated_data_rows = [
        *range(61, 66),
        *range(71, 76),
        *range(81, 86),
        *range(91, 96),
        *range(101, 106),
        *range(111, 116),
    ]
    repeated_height = max(float(ws.row_dimensions[row].height or 18.0) for row in repeated_data_rows)
    for row in repeated_data_rows:
        ws.row_dimensions[row].height = repeated_height
    subtotal_rows = (66, 76, 86, 96, 106, 116)
    subtotal_height = max(float(ws.row_dimensions[row].height or 18.0) for row in subtotal_rows)
    for row in subtotal_rows:
        ws.row_dimensions[row].height = subtotal_height

    # 空白分隔行保持紧凑。
    for row in (21, 23, 32, 39, 47, 55, 57, 58, 67, 68, 77, 78, 87, 88, 97, 98, 107, 108, 117, 118):
        ws.row_dimensions[row].height = 8.0


def configure_printing(ws) -> None:
    """固定 A4 纵向、一页宽，并只在完整表格之间分页。"""

    # H 列既承载“施工/供应说明”，也承载汇总页的金额差异。
    # 原始 12 字符列宽会让 1,310,760.00 在打印预览中显示为 ###。
    ws.column_dimensions["H"].width = 16.0

    ws.print_area = "A1:J124"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.scale = None
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.pageSetUpPr.autoPageBreaks = False
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4
    ws.page_margins.header = 0.2
    ws.page_margins.footer = 0.2

    ws.row_breaks = RowBreak()
    for row in PAGE_BREAK_ROWS:
        ws.row_breaks.append(Break(id=row))


def format_fees_sheet(ws) -> None:
    """只修正辅助表的表头可读性，不改变任何原始数据。"""

    ws.column_dimensions["C"].width = max(float(ws.column_dimensions["C"].width or 0), 18.0)
    ws.row_dimensions[1].height = max(float(ws.row_dimensions[1].height or 0), 30.0)
    for cell in ws[1]:
        cell.alignment = copy(cell.alignment)
        cell.alignment = Alignment(
            horizontal=cell.alignment.horizontal,
            vertical="center",
            text_rotation=cell.alignment.text_rotation,
            wrap_text=True,
            shrink_to_fit=cell.alignment.shrink_to_fit,
            indent=cell.alignment.indent,
        )


def solve(input_dir: Path, output_dir: Path) -> Path:
    input_path = input_dir / INPUT_NAME
    if not input_path.is_file():
        raise FileNotFoundError(f"缺少输入文件：{input_path}")

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / OUTPUT_NAME

    workbook = load_workbook(input_path, data_only=False)
    if MAIN_SHEET not in workbook.sheetnames:
        raise ValueError(f"输入工作簿缺少工作表：{MAIN_SHEET}")
    worksheet = workbook[MAIN_SHEET]

    # 截图只用于提供事实；输出必须使用真正的可编辑单元格。
    worksheet._images = []

    write_equipment_table(
        worksheet,
        title_row=69,
        title="五、环控与照明物资清单",
        subtotal_title="五、环控与照明物资清单 小计",
        records=ENVIRONMENT_RECORDS,
    )
    write_equipment_table(
        worksheet,
        title_row=79,
        title="六、物联网设备预算清单",
        subtotal_title="六、物联网设备预算清单 小计",
        records=IOT_RECORDS,
    )

    apply_uniform_table_typography(worksheet)
    adjust_row_heights(worksheet)
    configure_printing(worksheet)
    if "费用类别" in workbook.sheetnames:
        format_fees_sheet(workbook["费用类别"])

    # 让 Excel、WPS 或 LibreOffice 打开文件时重新计算所有新增公式。
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    print(solve(args.input_dir, args.output_dir))


if __name__ == "__main__":
    main()
