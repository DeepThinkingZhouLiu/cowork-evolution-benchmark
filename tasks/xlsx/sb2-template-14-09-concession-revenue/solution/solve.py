#!/usr/bin/env python3
"""Complete the visible revenue projection template with native formulas."""
from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

from openpyxl import load_workbook

INPUT_NAME = "concession_input.xlsx"
OUTPUT_NAME = "concession_revenue.xlsx"
EXPECTED_INPUT_SHA256 = "b272c7de3b687ecea8013d95e57cdeeccb071c3e7382d1a70abdff7b4216ec69"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def formulas() -> dict[str, str]:
    years = "DEFGH"
    expected: dict[str, str] = {
        "D32": "=C7",
        "D38": "=C17",
        "D39": "=C19",
        "D40": "=C21",
        "D41": "=C23",
    }
    for previous, current in zip(years, years[1:]):
        expected[f"{current}32"] = f"={previous}32*(1+$C$11)"
        expected[f"{current}38"] = f"={previous}38*(1+$C$18)"
        expected[f"{current}39"] = f"={previous}39*(1+$C$20)"
        expected[f"{current}40"] = f"={previous}40*(1+$C$22)"
        expected[f"{current}41"] = f"={previous}41*(1+$C$24)"
    for col in years:
        expected[f"{col}33"] = f"=-{col}32*$C$12"
    expected.update({"D34": "=-D32*$C$13", "E34": "=-E32*$C$14", "F34": "=0", "G34": "=0", "H34": "=0"})
    for col in years:
        expected[f"{col}35"] = f"=SUM({col}32:{col}34)"
        expected[f"{col}42"] = f"=SUM({col}38:{col}41)"
        expected[f"{col}44"] = f"={col}35+{col}42"
        expected[f"{col}45"] = f"=-{col}44*$C$27"
        expected[f"{col}46"] = f"={col}44+{col}45"
    for row in (32, 33, 34, 35, 38, 39, 40, 41, 42, 44, 45, 46):
        expected[f"I{row}"] = f"=SUM(D{row}:H{row})"
    return expected


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    source = args.input_dir / INPUT_NAME
    if sha256(source) != EXPECTED_INPUT_SHA256:
        raise ValueError("concession input hash mismatch")
    workbook = load_workbook(source)
    if workbook.sheetnames != ["Revenue_Projection"]:
        raise ValueError(f"unexpected workbook sheets: {workbook.sheetnames}")
    sheet = workbook["Revenue_Projection"]
    for coordinate, formula in formulas().items():
        sheet[coordinate] = formula
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    args.output_dir.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output_dir / OUTPUT_NAME)
    workbook.close()


if __name__ == "__main__":
    main()
