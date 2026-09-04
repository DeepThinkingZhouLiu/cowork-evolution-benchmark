#!/usr/bin/env python3
"""Independent deterministic Judge for the concession projection workbook."""
from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

TASK_ID = "cowork-evo/sb2-template-14-09-concession-revenue"
RUBRIC_VERSION = "cowork-evo-rubric-v0.2"
JUDGE_VERSION = "sb2-template-14-09-concession-revenue-deterministic-v0.2"
INPUT = "concession_input.xlsx"
OUTPUT = "concession_revenue.xlsx"
EXPECTED_INPUT_SHA256 = "b272c7de3b687ecea8013d95e57cdeeccb071c3e7382d1a70abdff7b4216ec69"
WEIGHTS = {"R001": 0.10, "R002": 0.45, "R003": 0.15, "R004": 0.20, "R005": 0.05, "R006": 0.05, "R007": -0.20}


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def criterion(cid: str, ok: bool, evidence: str) -> dict:
    return {"criterion_id": cid, "score": float(bool(ok)), "evidence": evidence, "method": "deterministic"}


def normalize(value):
    if hasattr(value, "text") and isinstance(value.text, str):
        value = value.text
    if isinstance(value, str) and value.startswith("="):
        return value.replace("$", "").replace(" ", "").upper()
    return value


def formula_text(value) -> str | None:
    if hasattr(value, "text") and isinstance(value.text, str):
        value = value.text
    return value if isinstance(value, str) and value.startswith("=") else None


def equal(left, right) -> bool:
    left, right = normalize(left), normalize(right)
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return abs(float(left) - float(right)) <= 1e-8
    return left == right or (left in (None, "") and right in (None, ""))


def row_by_label(sheet, label: str) -> int:
    for row in range(1, sheet.max_row + 1):
        if sheet.cell(row, 2).value == label:
            return row
    raise ValueError(f"missing template label: {label}")


def assumption_cell(sheet, label: str) -> str:
    return f"C{row_by_label(sheet, label)}"


def expected_formulas(source_sheet) -> dict[str, str]:
    years = [get_column_letter(col) for col in range(4, 9)]
    market = row_by_label(source_sheet, "Market Rent")
    loss = row_by_label(source_sheet, "Loss-to-Lease")
    concessions = row_by_label(source_sheet, "Concessions")
    rental_total = row_by_label(source_sheet, "Total Rental Revenue")
    rubs = row_by_label(source_sheet, "RUBS")
    amenity = row_by_label(source_sheet, "Amenity Income")
    storage = row_by_label(source_sheet, "Storage Income")
    parking = row_by_label(source_sheet, "Parking Income")
    other_total = row_by_label(source_sheet, "Total Other Income")
    gross = row_by_label(source_sheet, "GROSS REVENUE")
    vacancy = row_by_label(source_sheet, "General Vacancy & Credit Loss")
    effective = row_by_label(source_sheet, "EFFECTIVE GROSS REVENUE")
    result = {
        f"D{market}": f"={assumption_cell(source_sheet, 'Current Market Rent (Annual)')}",
        f"D{rubs}": f"={assumption_cell(source_sheet, 'RUBS - Year 1 (Annual)')}",
        f"D{amenity}": f"={assumption_cell(source_sheet, 'Amenity Income - Year 1 (Annual)')}",
        f"D{storage}": f"={assumption_cell(source_sheet, 'Storage Income - Year 1 (Annual)')}",
        f"D{parking}": f"={assumption_cell(source_sheet, 'Parking Income - Year 1 (Annual)')}",
        f"D{concessions}": f"=-D{market}*{assumption_cell(source_sheet, 'Year 1 Concessions (% of Market Rent)')}",
        f"E{concessions}": f"=-E{market}*{assumption_cell(source_sheet, 'Year 2 Concessions (% of Market Rent)')}",
    }
    growth = {
        market: "Market Rent Growth Rate",
        rubs: "RUBS Growth Rate",
        amenity: "Amenity Income Growth Rate",
        storage: "Storage Income Growth Rate",
        parking: "Parking Income Growth Rate",
    }
    for previous, current in zip(years, years[1:]):
        for row, assumption in growth.items():
            result[f"{current}{row}"] = f"={previous}{row}*(1+{assumption_cell(source_sheet, assumption)})"
    for col in years:
        result[f"{col}{loss}"] = f"=-{col}{market}*{assumption_cell(source_sheet, 'Loss-to-Lease (% of Market Rent)')}"
        result[f"{col}{rental_total}"] = f"=SUM({col}{market}:{col}{concessions})"
        result[f"{col}{other_total}"] = f"=SUM({col}{rubs}:{col}{parking})"
        result[f"{col}{gross}"] = f"={col}{rental_total}+{col}{other_total}"
        result[f"{col}{vacancy}"] = f"=-{col}{gross}*{assumption_cell(source_sheet, 'General Vacancy Rate (% of Gross Revenue)')}"
        result[f"{col}{effective}"] = f"={col}{gross}+{col}{vacancy}"
    for col in "FGH":
        result[f"{col}{concessions}"] = "=0"
    for row in (market, loss, concessions, rental_total, rubs, amenity, storage, parking, other_total, gross, vacancy, effective):
        result[f"I{row}"] = f"=SUM(D{row}:H{row})"
    return result


def numeric_truth(source_sheet) -> dict[str, float]:
    """Compute the expected values independently of any Golden workbook."""
    val = lambda label: float(source_sheet[assumption_cell(source_sheet, label)].value)
    market = val("Current Market Rent (Annual)")
    loss_pct = val("Loss-to-Lease (% of Market Rent)")
    market_growth = val("Market Rent Growth Rate")
    concession_rates = [val("Year 1 Concessions (% of Market Rent)"), val("Year 2 Concessions (% of Market Rent)"), 0.0, 0.0, 0.0]
    other = [
        ("RUBS", val("RUBS - Year 1 (Annual)"), val("RUBS Growth Rate")),
        ("Amenity Income", val("Amenity Income - Year 1 (Annual)"), val("Amenity Income Growth Rate")),
        ("Storage Income", val("Storage Income - Year 1 (Annual)"), val("Storage Income Growth Rate")),
        ("Parking Income", val("Parking Income - Year 1 (Annual)"), val("Parking Income Growth Rate")),
    ]
    vacancy_rate = val("General Vacancy Rate (% of Gross Revenue)")
    truth: dict[str, float] = {}
    for index, col in enumerate("DEFGH"):
        if index:
            market *= 1 + market_growth
        truth[f"{col}{row_by_label(source_sheet, 'Market Rent')}"] = market
        truth[f"{col}{row_by_label(source_sheet, 'Loss-to-Lease')}"] = -market * loss_pct
        truth[f"{col}{row_by_label(source_sheet, 'Concessions')}"] = -market * concession_rates[index]
        rental = sum(truth[f"{col}{row_by_label(source_sheet, label)}"] for label in ("Market Rent", "Loss-to-Lease", "Concessions"))
        total_other = 0.0
        for label, first, growth_rate in other:
            value = first * (1 + growth_rate) ** index
            truth[f"{col}{row_by_label(source_sheet, label)}"] = value
            total_other += value
        truth[f"{col}{row_by_label(source_sheet, 'Total Rental Revenue')}"] = rental
        truth[f"{col}{row_by_label(source_sheet, 'Total Other Income')}"] = total_other
        gross = rental + total_other
        truth[f"{col}{row_by_label(source_sheet, 'GROSS REVENUE')}"] = gross
        truth[f"{col}{row_by_label(source_sheet, 'General Vacancy & Credit Loss')}"] = -gross * vacancy_rate
        truth[f"{col}{row_by_label(source_sheet, 'EFFECTIVE GROSS REVENUE')}"] = gross * (1 - vacancy_rate)
    for row in ("Market Rent", "Loss-to-Lease", "Concessions", "Total Rental Revenue", "RUBS", "Amenity Income", "Storage Income", "Parking Income", "Total Other Income", "GROSS REVENUE", "General Vacancy & Credit Loss", "EFFECTIVE GROSS REVENUE"):
        row_number = row_by_label(source_sheet, row)
        truth[f"I{row_number}"] = sum(truth[f"{col}{row_number}"] for col in "DEFGH")
    return truth


def recalc_values(path: Path, coordinates: list[str]) -> dict[str, float] | None:
    """Use LibreOffice when present; static formula checks remain authoritative."""
    binary = shutil.which("libreoffice") or shutil.which("soffice")
    if not binary:
        return None
    with tempfile.TemporaryDirectory() as temp:
        source_dir = Path(temp) / "source"
        output_dir = Path(temp) / "recalculated"
        source_dir.mkdir()
        output_dir.mkdir()
        temp_path = source_dir / path.name
        shutil.copy2(path, temp_path)
        try:
            subprocess.run([binary, "--headless", "--convert-to", "xlsx", "--outdir", str(output_dir), str(temp_path)], check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=60)
        except (OSError, subprocess.SubprocessError):
            return None
        try:
            recalculated_path = output_dir / path.name
            wb = load_workbook(recalculated_path, data_only=True)
            ws = wb["Revenue_Projection"]
            values = {coordinate: ws[coordinate].value for coordinate in coordinates}
            wb.close()
            return values
        except Exception:
            return None


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--reward-file", type=Path, required=True)
    args = parser.parse_args()
    root = Path(__file__).parents[1]
    source_path = root / "data" / "input_files" / INPUT
    output_path = args.artifact_dir / OUTPUT
    results: list[dict] = []
    native = output_path.is_file() and zipfile.is_zipfile(output_path)
    results.append(criterion("R001", native, "output exists and is a native OOXML workbook"))
    formula_ok = structure_ok = preserve_ok = calc_ok = False
    if native:
        try:
            source = load_workbook(source_path, data_only=False)
            output = load_workbook(output_path, data_only=False)
            source_sheet = source["Revenue_Projection"]
            structure_ok = output.sheetnames == ["Revenue_Projection"] and output.active.max_row == source_sheet.max_row and output.active.max_column == source_sheet.max_column
            output_sheet = output["Revenue_Projection"] if "Revenue_Projection" in output.sheetnames else output.active
            expected = expected_formulas(source_sheet)
            mismatches = [coordinate for coordinate, formula in expected.items() if normalize(output_sheet[coordinate].value) != normalize(formula)]
            formula_ok = structure_ok and not mismatches
            results.append(criterion("R002", formula_ok, f"expected formula cells={len(expected)}, mismatches={mismatches[:8]}"))
            truth = numeric_truth(source_sheet)
            recalculated = recalc_values(output_path, list(truth))
            if recalculated is None:
                numeric_ok = formula_ok
                numeric_evidence = "formula semantics checked; optional LibreOffice replay unavailable"
            else:
                numeric_mismatches = [coordinate for coordinate, value in truth.items() if not isinstance(recalculated.get(coordinate), (int, float)) or abs(float(recalculated[coordinate]) - value) > 1e-5 * max(1.0, abs(value))]
                numeric_ok = not numeric_mismatches
                numeric_evidence = f"recalculated mismatches={numeric_mismatches[:8]}"
            results.append(criterion("R003", numeric_ok, numeric_evidence))
            target_coordinates = set(expected)
            preserve_ok = True
            for row in range(1, source_sheet.max_row + 1):
                for col in range(1, source_sheet.max_column + 1):
                    coordinate = source_sheet.cell(row, col).coordinate
                    if coordinate not in target_coordinates and not equal(source_sheet.cell(row, col).value, output_sheet.cell(row, col).value):
                        preserve_ok = False
            preserve_ok = preserve_ok and output.sheetnames == source.sheetnames and output_sheet.merged_cells.ranges == source_sheet.merged_cells.ranges
            preserve_ok = preserve_ok and output_sheet.freeze_panes == source_sheet.freeze_panes and output_sheet.auto_filter.ref == source_sheet.auto_filter.ref and output_sheet.print_area == source_sheet.print_area
            results.append(criterion("R004", preserve_ok, "non-target cell values, sheet order, merges and view settings preserved"))
            formula_count = sum(1 for sheet in output.worksheets for row in sheet.iter_rows() for cell in row if formula_text(cell.value) is not None)
            broken = any("#REF!" in formula_text(cell.value).upper() for sheet in output.worksheets for row in sheet.iter_rows() for cell in row if formula_text(cell.value) is not None)
            calc = output.calculation
            calc_ok = formula_ok and formula_count == len(expected) and not broken and calc.calcMode == "auto" and bool(calc.fullCalcOnLoad) and bool(calc.forceFullCalc)
            results.append(criterion("R005", calc_ok, f"formula_count={formula_count}, broken_refs={broken}, calculation={calc.calcMode}/{calc.fullCalcOnLoad}/{calc.forceFullCalc}"))
            results.append(criterion("R006", digest(source_path) == EXPECTED_INPUT_SHA256, f"input_sha256={digest(source_path)}"))
            harmful = digest(output_path) == digest(source_path) or not formula_ok
            results.append(criterion("R007", harmful, "unchanged or hardcoded/incorrect output penalty" if harmful else "no unchanged or hardcoded substitution detected"))
            source.close()
            output.close()
        except Exception as exc:
            results.extend(criterion(cid, False, f"Judge failure: {exc}") for cid in ("R002", "R003", "R004", "R005", "R006"))
            results.append(criterion("R007", True, "invalid output penalty"))
    else:
        results.extend(criterion(cid, False, "missing or invalid output") for cid in ("R002", "R003", "R004", "R005"))
        results.append(criterion("R006", digest(source_path) == EXPECTED_INPUT_SHA256, "input hash"))
        results.append(criterion("R007", True, "missing or invalid output penalty"))
    values = {item["criterion_id"]: item["score"] for item in results}
    reward = sum(WEIGHTS[key] * values.get(key, 0.0) for key in WEIGHTS)
    passed = all(values.get(key, 0) == 1 for key in ("R001", "R002", "R004")) and reward >= 0.7
    payload = {"task_id": TASK_ID, "rubric_version": RUBRIC_VERSION, "judge_version": JUDGE_VERSION, "artifact_sha256": [digest(output_path)] if output_path.is_file() else [], "criterion_results": results, "reward": reward, "passed": passed}
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    args.reward_file.parent.mkdir(parents=True, exist_ok=True)
    args.reward_file.write_text(f"{reward:.6f}\n", encoding="utf-8")
    raise SystemExit(0 if passed else 1)


if __name__ == "__main__":
    main()
