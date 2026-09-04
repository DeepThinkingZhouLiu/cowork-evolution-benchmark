#!/usr/bin/env python3
from __future__ import annotations
import argparse, re
from datetime import timedelta
from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import column_index_from_string

KIND = 'archer'

def numeric(value):
    if value is None or value == "": return None
    if isinstance(value, timedelta): return value.total_seconds() / 3600.0
    if isinstance(value, (int, float)): return float(value) * 24.0 if 0 <= float(value) < 2 else float(value)
    text = str(value).replace("\xa0", " ").strip().replace(",", "")
    if not text or text in {"-", "--"}: return None
    if ":" in text:
        left, right = text.split(":", 1)
        try: return float(left.strip()) + float(right.strip()) / 60.0
        except ValueError: return None
    try: return float(text)
    except ValueError: return None

def rightmost_match(left, right, rank_col, row):
    start = column_index_from_string(left); end = column_index_from_string(right)
    parts = []
    for col in range(end, start - 1, -1):
        letter = get_column_letter(col)
        parts.append("IF(" + letter + str(row) + "=$" + rank_col + str(row) + "," + str(col) + ",")
    return "".join(parts) + '""' + ")" * len(parts)

def dropdown(src, dst):
    wb = load_workbook(src)
    from openpyxl.utils import get_column_letter
    form = wb["Form"]
    lists = wb.create_sheet("Lists", wb.sheetnames.index("Form"))
    categories = []
    grouped = {}
    ws = wb["Account"]
    for row in range(2, ws.max_row + 1):
        cat, sub = ws.cell(row, 3).value, ws.cell(row, 4).value
        if cat not in grouped: categories.append(cat); grouped[cat] = []
        if sub not in grouped[cat]: grouped[cat].append(sub)
    for col, cat in enumerate(categories, 1):
        lists.cell(1, col).value = cat
        for row, sub in enumerate(grouped[cat], 2): lists.cell(row, col).value = sub
        safe = re.sub(r"[^A-Za-z0-9_]", "_", str(cat).replace(" ", "_"))
        wb.defined_names.add(DefinedName(safe, attr_text="Lists!$" + get_column_letter(col) + "$2:$" + get_column_letter(col) + "$" + str(len(grouped[cat]) + 1)))
    wb.defined_names.add(DefinedName("Categories", attr_text="Lists!$A$1:$" + get_column_letter(len(categories)) + "$1"))
    for dv in list(form.data_validations.dataValidation): form.data_validations.dataValidation.remove(dv)
    dv1 = DataValidation(type="list", formula1="Categories"); dv1.add(form["B2"]); form.add_data_validation(dv1)
    dv2 = DataValidation(type="list", formula1='INDIRECT(SUBSTITUTE(B2, " ", "_"))'); dv2.add(form["C2"]); form.add_data_validation(dv2)
    for col in range(1, len(categories)+1): lists.column_dimensions[get_column_letter(col)].width = 22
    wb.calculation.calcMode = "auto"; wb.calculation.fullCalcOnLoad = True; wb.calculation.forceFullCalc = True
    wb.save(dst)

def reshape(src, dst):
    source = load_workbook(src, data_only=True)["DB"]
    wb = Workbook(); ws = wb.active; ws.title = "Sheet1"
    areas = []
    category_area = []
    for col in range(2, source.max_column + 1):
        cat = str(source.cell(4, col).value or "")
        cat = re.sub(r"\d+$", "", cat)
        area = source.cell(5, col).value
        if area not in areas: areas.append(area)
        category_area.append((cat, area, col))
    cats = sorted(dict.fromkeys(item[0] for item in category_area))
    headers = ["Removed Characters", "Attribute"] + areas
    for col, value in enumerate(headers, 1): ws.cell(1, col).value = value
    row = 2
    for cat in cats:
        for seq in range(1, 6):
            ws.cell(row, 1).value = cat; ws.cell(row, 2).value = seq
            for out_col, area in enumerate(areas, 3):
                src_col = next((c for ccat, carea, c in category_area if ccat == cat and carea == area), None)
                if src_col is not None: ws.cell(row, out_col).value = numeric(source.cell(seq + 5, src_col).value)
            row += 1
    for cell in ws[1]: cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="D9EAF7")
    ws.freeze_panes = "C2"; ws.auto_filter.ref = "A1:" + get_column_letter(len(headers)) + str(row - 1)
    wb.save(dst)

def hours(src, dst):
    wb = load_workbook(src)
    data = wb["Data"]
    if "Analysis" in wb.sheetnames: del wb["Analysis"]
    ws = wb.create_sheet("Analysis")
    # Data column AE is the source TOTAL; only C:AD are the 28 month inputs.
    from openpyxl.utils import get_column_letter
    month_cols = range(3, data.max_column)
    months = [data.cell(1, c).value for c in month_cols]
    for c, value in enumerate(["Department"] + months + ["TOTAL"], 1): ws.cell(1, c).value = value
    for r in range(3, 25):
        out_r = r - 1; ws.cell(out_r, 1).value = f"=Data!B{r}"
        for c in month_cols:
            out_c = c - 1; source = f"Data!{__import__('openpyxl').utils.get_column_letter(c)}{r}"
            ws.cell(out_r, out_c).value = f'=IF(ISNUMBER({source}),{source}*24,IFERROR(VALUE(SUBSTITUTE(LEFT({source},SEARCH(":",{source})-1),",","")),0)+IFERROR(VALUE(RIGHT({source},LEN({source})-SEARCH(":",{source})))/60,0))'
        last = get_column_letter(data.max_column - 2)
        ws.cell(out_r, data.max_column - 1).value = f"=SUM(B{out_r}:{last}{out_r})"
    ws.cell(25, 1).value = "TOTAL"
    for c in range(2, data.max_column):
        col = get_column_letter(c); ws.cell(25, c).value = f"=SUM({col}2:{col}23)"
    ws.cell(25, data.max_column - 1).value = "=SUM(AD2:AD23)"
    chart = LineChart(); chart.title = "Monthly Hours per Department"; chart.y_axis.title = "Hours"; chart.x_axis.title = "Month"
    chart.add_data(Reference(ws, min_col=2, max_col=data.max_column-2, min_row=1, max_row=23), titles_from_data=False, from_rows=True)
    chart.set_categories(Reference(ws, min_col=2, max_col=data.max_column-2, min_row=1, max_row=1)); chart.height = 12; chart.width = 25; chart.anchor = "B27"; ws.add_chart(chart)
    for cell in ws[1]: cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="D9EAF7")
    ws.freeze_panes = "B2"; ws.auto_filter.ref = "A1:" + get_column_letter(data.max_column - 1) + "23"
    wb.calculation.calcMode = "auto"; wb.calculation.fullCalcOnLoad = True; wb.calculation.forceFullCalc = True; wb.save(dst)

def outer(src, dst):
    wb = load_workbook(src); t1, t2 = wb["Table1"], wb["Table2"]
    if "Output" in wb.sheetnames: del wb["Output"]
    ws = wb.create_sheet("Output")
    headers = ["KEY CODE", "PRODUCT NAME (T1)", "PRODUCT CODE (T1)", "SALES (T1)", "COGS (T1)", "PROFIT (T1)", "PRODUCT NAME (T2)", "PRODUCT CODE (T2)", "SALES (T2)", "COGS (T2)", "PROFIT (T2)", "COMBINED PRODUCT NAME", "COMBINED PRODUCT CODE"]
    for c, value in enumerate(headers, 1): ws.cell(1, c).value = value
    keys = []
    for source in (t1, t2):
        for r in range(2, source.max_row + 1):
            code = source.cell(r, 2).value
            if code not in keys: keys.append(code)
    from openpyxl.utils import get_column_letter
    for r, key in enumerate(keys, 2):
        ws.cell(r, 1).value = key
        for c in range(2, 12):
            source_sheet = "Table1" if c <= 6 else "Table2"; source_col = c - 1 if c <= 6 else c - 6
            ws.cell(r, c).value = f'=IFERROR(INDEX({source_sheet}!$A:$E,MATCH($A{r},{source_sheet}!$B:$B,0),{source_col}),"")'
        ws.cell(r, 12).value = f'=IF($B{r}<>"",$B{r},$G{r})'; ws.cell(r, 13).value = f'=IF($C{r}<>"",$C{r},$H{r})'
    for cell in ws[1]: cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="D9EAF7")
    ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:M{len(keys)+1}"; wb.calculation.calcMode = "auto"; wb.calculation.fullCalcOnLoad = True; wb.calculation.forceFullCalc = True; wb.save(dst)

def archer(src, dst):
    wb = load_workbook(src); ws = wb["NEW"]
    from openpyxl.utils import get_column_letter
    groups = [("C","D",4),("F","G",7),("I","J",10),("L","M",13),("O","Q",16),("S","T",19)]
    for r in range(3, 33):
        maxes = [f"MAX(${a}{r}:${b}{r})" for a,b,_ in groups]
        for offset, col in enumerate((25,26,27,28), 1): ws.cell(r, col).value = f'=LARGE(CHOOSE({{1,2,3,4,5,6}},{",".join(maxes)}),{offset})'
        for j, col in enumerate((30,31,32,33), 0):
            tests = []
            for a,b,right_col in groups:
                used = "TRUE" if j == 0 else f"COUNTIF($AD{r}:${get_column_letter(29+j)}{r},{right_col})=0"
                rank_col = get_column_letter(25 + j)
                match = rightmost_match(a, b, rank_col, r)
                tests.append(f'IF(AND(${rank_col}{r}<>0,MAX(${a}{r}:${b}{r})=${rank_col}{r},{used}),{match},')
            formula = "=" + "".join(tests) + '""' + ")" * len(tests)
            ws.cell(r, col).value = formula
    strike = Font(strike=True, color="9C0006"); ws.conditional_formatting.add("C3:T32", FormulaRule(formula=["=AND(C3<>"",COUNTIF($Y3:$AB3,C3)<COUNTIF($C3:$T3,C3))"], font=strike))
    ws.column_dimensions["Y"].width = 12; ws.column_dimensions["Z"].width = 12; ws.column_dimensions["AA"].width = 12; ws.column_dimensions["AB"].width = 12; ws.column_dimensions["AD"].width = 10; ws.column_dimensions["AE"].width = 10; ws.column_dimensions["AF"].width = 10; ws.column_dimensions["AG"].width = 10
    wb.calculation.calcMode = "auto"; wb.calculation.fullCalcOnLoad = True; wb.calculation.forceFullCalc = True; wb.save(dst)

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--input-dir", type=Path, required=True); ap.add_argument("--output-dir", type=Path, required=True); args = ap.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True); src = args.input_dir / "archer_scores_input.xlsx"
    out = args.output_dir / "archer_scores_ranked.xlsx"
    {"dropdown": dropdown, "reshape": reshape, "hours": hours, "outer": outer, "archer": archer}[KIND](src, out)
if __name__ == "__main__": main()
