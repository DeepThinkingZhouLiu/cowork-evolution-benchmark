#!/usr/bin/env python3
from __future__ import annotations
import argparse
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

INPUT_NAME="transactions_input.xlsx"; OUTPUT_NAME="transactions_fifo.xlsx"
def solve(input_dir: Path, output_dir: Path) -> None:
    wb=load_workbook(input_dir/INPUT_NAME,data_only=False); tx=wb["Transactions"]
    buys=[]; sells=[]
    for r in range(2,tx.max_row+1):
        typ=tx.cell(r,4).value
        if typ=="BUY": buys.append(r)
        elif typ=="SELL": sells.append(r)
    if "FIFO_Alloc" in wb.sheetnames: del wb["FIFO_Alloc"]
    ws=wb.create_sheet("FIFO_Alloc")
    ws["B1"]="Sale Date"; ws["C1"]="Sell Qty"; ws["AA1"]="Relieved Cost ($)"; ws["AB1"]="Avg Cost (c/$)"
    ws["A2"]="Buy Qty"; ws["A3"]="Buy Unit Cost ($/$)"
    for idx,r in enumerate(buys,start=4):
        col=get_column_letter(idx)
        ws.cell(1,idx,f"=Transactions!E{r}"); ws.cell(2,idx,f"=Transactions!F{r}"); ws.cell(3,idx,f"=Transactions!G{r}/100")
        ws.cell(1,idx).number_format="yyyy-mm-dd"; ws.cell(3,idx).number_format="0.0000"
    for out_row,tx_row in enumerate(sells,start=4):
        ws.cell(out_row,1,f"Sale r{tx_row}"); ws.cell(out_row,2,f"=Transactions!E{tx_row}"); ws.cell(out_row,3,f"=Transactions!F{tx_row}")
        ws.cell(out_row,2).number_format="yyyy-mm-dd"
        for idx in range(4,4+len(buys)):
            col=get_column_letter(idx); prior=f"{col}$4:{col}{out_row-1}" if out_row>4 else None; remaining=f"{col}$2" if prior is None else f"{col}$2-SUM({prior})"; sold=f"SUM($D{out_row}:{get_column_letter(idx-1)}{out_row})" if idx>4 else "0"
            ws.cell(out_row,idx,f"=IF({col}$1<=$B{out_row},MAX(0,MIN({remaining},$C{out_row}-({sold}))),0)")
        end=get_column_letter(3+len(buys)); ws.cell(out_row,27,f"=SUMPRODUCT(D{out_row}:{end}{out_row},D$3:{end}$3)"); ws.cell(out_row,28,f"=IF($C{out_row}>0,AA{out_row}/$C{out_row}*100,\"\")")
    for cell in ws[1]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="1F4E78"); cell.alignment=Alignment(horizontal="center")
    for cell in ws[2]: cell.font=Font(bold=True)
    ws.freeze_panes="D4"; ws.column_dimensions["A"].width=16; ws.column_dimensions["B"].width=14; ws.column_dimensions["C"].width=14; ws.column_dimensions["AA"].width=18; ws.column_dimensions["AB"].width=18
    for col in range(4,4+len(buys)): ws.column_dimensions[get_column_letter(col)].width=13
    wb.calculation.fullCalcOnLoad=True; wb.calculation.forceFullCalc=True; wb.calculation.calcMode="auto"
    output_dir.mkdir(parents=True,exist_ok=True); wb.save(output_dir/OUTPUT_NAME); wb.close()
def main():
    ap=argparse.ArgumentParser(); ap.add_argument("--input-dir",type=Path,required=True); ap.add_argument("--output-dir",type=Path,required=True); a=ap.parse_args(); solve(a.input_dir,a.output_dir)
if __name__=="__main__": main()
