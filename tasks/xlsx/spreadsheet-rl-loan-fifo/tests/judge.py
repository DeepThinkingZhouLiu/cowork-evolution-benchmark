#!/usr/bin/env python3
from __future__ import annotations
import argparse, hashlib, json, shutil, subprocess, tempfile, zipfile, io
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
OUT="transactions_fifo.xlsx"; TASK_ID="cowork-evo/spreadsheet-rl-loan-fifo"; WEIGHTS={"R001":.15,"R002":.20,"R003":.30,"R004":.20,"R005":.10,"R006":.05}
def crit(cid,ok,ev): return {"criterion_id":cid,"score":1.0 if ok else 0.0,"evidence":ev,"method":"deterministic"}
def recalc(path):
    lo=shutil.which("libreoffice") or shutil.which("soffice")
    if not lo:return None
    with tempfile.TemporaryDirectory(prefix="fifo-") as td:
        r=subprocess.run([lo,"--headless","--convert-to","xlsx","--outdir",td,str(path)],capture_output=True,text=True,timeout=180,check=False); p=Path(td)/path.name
        return load_workbook(io.BytesIO(p.read_bytes()),data_only=True) if r.returncode==0 and p.is_file() else None
def main():
    ap=argparse.ArgumentParser(); ap.add_argument("--artifact-dir",type=Path,required=True); ap.add_argument("--result",type=Path,required=True); ap.add_argument("--reward-file",type=Path,required=True); a=ap.parse_args(); source=Path(__file__).parents[1]/"data/input_files/transactions_input.xlsx"; out=a.artifact_dir/OUT; results=[]
    try:
        native=out.is_file() and zipfile.is_zipfile(out); results.append(crit("R001",native,"transactions_fifo.xlsx 可打开"))
        src=load_workbook(source,data_only=True,read_only=True)["Transactions"]; buys=[]; sells=[]
        for r in range(2,src.max_row+1):
            typ=src.cell(r,4).value
            if typ=="BUY": buys.append([src.cell(r,6).value,src.cell(r,7).value,r])
            elif typ=="SELL": sells.append([src.cell(r,6).value,r])
        expected=[]
        for qty,_ in sells:
            rem=qty; cost=0
            for lot in buys:
                take=min(rem,lot[0]); cost+=take*lot[1]; lot[0]-=take; rem-=take
                if rem<=0:break
            expected.append(cost/qty*1.0)
        formulas_ok=False; structure=False
        if native:
            wf=load_workbook(out,data_only=False); ws=wf["FIFO_Alloc"] if "FIFO_Alloc" in wf.sheetnames else None; structure=bool(ws and tuple(ws.cell(1,c).value for c in [2,3,27,28])==("Sale Date","Sell Qty","Relieved Cost ($)","Avg Cost (c/$)")); formulas_ok=bool(ws and all(isinstance(ws.cell(4+i,28).value,str) and ws.cell(4+i,28).value.startswith("=") for i in range(len(sells)))); results.append(crit("R002",structure,"FIFO_Alloc 结构和四个关键列存在")); results.append(crit("R003",formulas_ok,"每笔 SELL 的 AB 结果为公式"));
            preserved=wf.sheetnames[0]=="Transactions" and wf["Transactions"].max_row==src.max_row and wf["Transactions"].max_column==src.max_column
            results.append(crit("R005",preserved,"原 Transactions 工作表保留")); wf.close()
        else: results.extend([crit("R002",False,"缺少 FIFO_Alloc"),crit("R003",False,"缺少公式"),crit("R005",False,"无法检查原表")])
        calc=recalc(out) if native else None; values_ok=False
        if calc and "FIFO_Alloc" in calc.sheetnames:
            ws=calc["FIFO_Alloc"]; actual=[ws.cell(4+i,28).value for i in range(len(sells))]; values_ok=all(v is not None and abs(float(v)-e)<1e-6 for v,e in zip(actual,expected)); calc.close()
        results.append(crit("R004",values_ok,f"FIFO 期望={expected}")); results.append(crit("R006",hashlib.sha256(source.read_bytes()).hexdigest()=="3195814da1d869b35721a1a60f100192b5ddd2e106971c9856a0b0a33def2fa2","输入 SHA256 固定"))
    except Exception as exc: results=[crit("R001",False,f"Judge 读取失败：{exc}")]
    vals={x["criterion_id"]:x["score"] for x in results}; reward=sum(WEIGHTS[k]*vals.get(k,0) for k in WEIGHTS); passed=reward>=.8 and vals.get("R001",0)==1 and vals.get("R003",0)==1 and vals.get("R004",0)==1
    payload={"task_id":TASK_ID,"rubric_version":"spreadsheet-rl-loan-fifo-v0.1","judge_version":"spreadsheet-rl-loan-fifo-deterministic-v0.1","artifact_sha256":[hashlib.sha256(out.read_bytes()).hexdigest()] if out.is_file() else [],"criterion_results":results,"reward":reward,"passed":passed}; a.result.parent.mkdir(parents=True,exist_ok=True); a.result.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8"); a.reward_file.parent.mkdir(parents=True,exist_ok=True); a.reward_file.write_text(f"{reward:.6f}\n"); raise SystemExit(0 if passed else 1)
if __name__=="__main__": main()
