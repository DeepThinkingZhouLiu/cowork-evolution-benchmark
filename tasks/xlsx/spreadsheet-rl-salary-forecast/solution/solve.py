#!/usr/bin/env python3
from __future__ import annotations
import argparse
from pathlib import Path
from openpyxl import load_workbook

INPUT_NAME="salary_model_input.xlsx"; OUTPUT_NAME="salary_model_completed.xlsx"
def solve(input_dir: Path, output_dir: Path) -> None:
    wb=load_workbook(input_dir/INPUT_NAME,data_only=False); ws=wb["P10 SFT"]
    for r in range(8,16):
        ws[f"M{r}"] = f"=IF(H{r}=\"\",E{r},IFERROR(E{r}*DAY(H{r}-1)/DAY(EOMONTH(H{r},0))+L{r}*(DAY(EOMONTH(H{r},0))-DAY(H{r}-1))/DAY(EOMONTH(H{r},0)),E{r}))"
        ws[f"O{r}"] = f"=MAX($D{r},IFERROR(INDEX(CBAIncrement[[Step 1]:[Step 9]],MATCH($F{r},CBAIncrement[Level],0),$G{r}+IF($H{r}<$B$34,1,0)),$D{r}))"
        ws[f"Q{r}"] = f"=IF($H{r}<$B$21,$L{r},$E{r})*21/31+$P{r}*10/31"
    wb.calculation.fullCalcOnLoad=True; wb.calculation.forceFullCalc=True; wb.calculation.calcMode="auto"
    output_dir.mkdir(parents=True,exist_ok=True); wb.save(output_dir/OUTPUT_NAME); wb.close()
def main():
    ap=argparse.ArgumentParser(); ap.add_argument("--input-dir",type=Path,required=True); ap.add_argument("--output-dir",type=Path,required=True); a=ap.parse_args(); solve(a.input_dir,a.output_dir)
if __name__=="__main__": main()
