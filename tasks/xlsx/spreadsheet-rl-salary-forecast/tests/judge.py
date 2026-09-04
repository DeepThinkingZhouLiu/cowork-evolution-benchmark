#!/usr/bin/env python3
from __future__ import annotations
import argparse, hashlib, json, math, zipfile, shutil, subprocess, tempfile, io
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook
OUT="salary_model_completed.xlsx"; TASK_ID="cowork-evo/spreadsheet-rl-salary-forecast"; WEIGHTS={"R001":.15,"R002":.25,"R003":.30,"R004":.15,"R005":.10,"R006":.05}
def crit(cid,ok,ev): return {"criterion_id":cid,"score":1.0 if ok else 0.0,"evidence":ev,"method":"deterministic"}
def month_days(dt):
    if dt.month==12: nxt=datetime(dt.year+1,1,1)
    else: nxt=datetime(dt.year,dt.month+1,1)
    return (nxt-datetime(dt.year,dt.month,1)).days
def recalc(path: Path):
    lo=shutil.which("libreoffice") or shutil.which("soffice")
    if not lo: return None
    with tempfile.TemporaryDirectory(prefix="salary-") as td:
        r=subprocess.run([lo,"--headless","--convert-to","xlsx","--outdir",td,str(path)],capture_output=True,text=True,timeout=240,check=False)
        p=Path(td)/path.name
        return load_workbook(io.BytesIO(p.read_bytes()),data_only=True) if r.returncode==0 and p.is_file() else None
def main():
    ap=argparse.ArgumentParser(); ap.add_argument("--artifact-dir",type=Path,required=True); ap.add_argument("--result",type=Path,required=True); ap.add_argument("--reward-file",type=Path,required=True); a=ap.parse_args(); source=Path(__file__).parents[1]/"data/input_files/salary_model_input.xlsx"; out=a.artifact_dir/OUT; results=[]
    try:
        wi=load_workbook(source,data_only=True); si=wi["P10 SFT"]; ann={}; cba={}
        for r in range(21,31):
            level=si.cell(r,1).value
            if level: ann[level]=[si.cell(r,c).value for c in range(3,12)]
        for r in range(34,44):
            level=si.cell(r,1).value
            if level: cba[level]=[si.cell(r,c).value for c in range(3,12)]
        expected={}
        cba_date=si["B34"].value; ann_date=si["B21"].value
        for r in range(8,16):
            D,E,F,G,H,K,L=si[f"D{r}"].value,si[f"E{r}"].value,si[f"F{r}"].value,si[f"G{r}"].value,si[f"H{r}"].value,si[f"K{r}"].value,si[f"L{r}"].value
            if H:
                days=month_days(H); M=E*(H.day-1)/days+L*(days-H.day+1)/days
            else: M=E
            idx=int(G)-1+(1 if (H is None or H<cba_date) else 0); row=cba.get(F,[]); O=max(float(D),float(row[idx]) if 0<=idx<len(row) and row[idx] is not None else float(D)); P=O/12; base=L if H is None or H<ann_date else E; Q=base*21/31+P*10/31
            expected[r]=(M,O,Q)
        native=out.is_file() and zipfile.is_zipfile(out); results.append(crit("R001",native,"输出为可打开 XLSX")); formulas_ok=False; preserved=False
        if native:
            wo=load_workbook(out,data_only=False); so=wo["P10 SFT"]; formulas_ok=all(isinstance(so[f"{c}{r}"].value,str) and so[f"{c}{r}"].value.startswith("=") for r in range(8,16) for c in ("M","O","Q")); preserved=wo.sheetnames==wi.sheetnames and so.max_row==si.max_row and so.max_column==si.max_column; results.append(crit("R002",formulas_ok,"M/O/Q 24 个目标单元格均为公式")); results.append(crit("R005",preserved,"工作表结构保留"));
            actual_wb=recalc(out)
            if actual_wb:
                ao=actual_wb["P10 SFT"]; values_ok=all(ao[f"{c}{r}"].value is not None and abs(float(ao[f"{c}{r}"].value)-expected[r][i])<1e-5 for r in expected for i,c in enumerate(("M","O","Q"))); actual_wb.close()
            else:
                values_ok=False
            wo.close()
        else: values_ok=False; results.extend([crit("R002",False,"缺少公式"),crit("R005",False,"无法检查结构")])
        results.append(crit("R003",values_ok,"从增幅表、日期和基准工资独立重算 M/O/Q")); results.append(crit("R004",all(len(expected[r])==3 for r in expected),"八行三列预期值可确定")); results.append(crit("R006",hashlib.sha256(source.read_bytes()).hexdigest()=="1e1befbca7c4ef01a6d1a0ff8145d6569b1867321f6d0bc5964e66c59e356457","输入 SHA256 固定")); wi.close()
    except Exception as exc: results=[crit("R001",False,f"Judge 读取失败：{exc}")]
    vals={x["criterion_id"]:x["score"] for x in results}; reward=sum(WEIGHTS[k]*vals.get(k,0) for k in WEIGHTS); passed=reward>=.8 and vals.get("R001",0)==1 and vals.get("R002",0)==1 and vals.get("R003",0)==1
    payload={"task_id":TASK_ID,"rubric_version":"spreadsheet-rl-salary-forecast-v0.1","judge_version":"spreadsheet-rl-salary-forecast-deterministic-v0.1","artifact_sha256":[hashlib.sha256(out.read_bytes()).hexdigest()] if out.is_file() else [],"criterion_results":results,"reward":reward,"passed":passed}; a.result.parent.mkdir(parents=True,exist_ok=True); a.result.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8"); a.reward_file.parent.mkdir(parents=True,exist_ok=True); a.reward_file.write_text(f"{reward:.6f}\n"); raise SystemExit(0 if passed else 1)
if __name__=="__main__": main()
