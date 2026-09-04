#!/usr/bin/env python3
"""独立检查单一 LBO 修复工作簿的公式与结构。"""
from __future__ import annotations
import argparse, hashlib, json, zipfile
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

TASK_ID="cowork-evo/spreadsheetbench2-09-01-lbo-formula-repair"
RUBRIC_VERSION="spreadsheetbench2-09-01-lbo-formula-repair-v0.2"
JUDGE_VERSION="spreadsheetbench2-09-01-lbo-formula-repair-deterministic-v0.2"
OUTPUT="09_01_repaired.xlsx"
EXPECTED_SHA="f8a833086c28d4fd21e1bd11a083e2e8bd022002f5e0f6c2c57ae4ff9df80a44"
RANGES={"LBO":(2,218,1,24),"Financials":(2,29,1,17),"Operating Cases":(2,47,1,17)}
WEIGHTS={"R001":.10,"R002":.30,"R003":.15,"R004":.15,"R005":.10,"R006":.10,"R007":.05,"R008":.05}

def digest(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def criterion(cid, ok, evidence): return {"criterion_id":cid,"score":float(ok),"evidence":evidence,"method":"deterministic"}
def norm(value):
    if isinstance(value,str) and value.startswith("="):
        value=value.replace("$","").replace(" ","").upper()
        return "="+value[2:] if value.startswith("=+") else value
    return value
def equal(a,b):
    a,b=norm(a),norm(b)
    if isinstance(a,(int,float)) and isinstance(b,(int,float)): return abs(float(a)-float(b))<1e-10
    return a==b or (a in (None,"") and b in (None,""))

def targets():
    result={}
    for n in range(7,15):
        col=get_column_letter(n)
        result[("LBO",f"{col}66")]=f"={col}64+{col}75"
        result[("LBO",f"{col}69")]=f"=SUM({col}66:{col}68)"
        result[("LBO",f"{col}175")]=f"={col}136+{col}137+{col}145+{col}152"
    for n in range(10,15):
        col=get_column_letter(n); result[("Financials",f"{col}18")]=f"={col}12-SUM({col}15:{col}17)"
    return result

def main():
    ap=argparse.ArgumentParser(); ap.add_argument("--artifact-dir",type=Path,required=True); ap.add_argument("--result",type=Path,required=True); ap.add_argument("--reward-file",type=Path,required=True); args=ap.parse_args()
    root=Path(__file__).parents[1]; source_path=root/"data/input_files/09_01_input.xlsx"; out_path=args.artifact_dir/OUTPUT; results=[]
    valid=out_path.is_file() and zipfile.is_zipfile(out_path)
    if not valid:
        results=[criterion("R001",False,"输出缺失或不是 OOXML")]+[criterion(cid,False,"缺少有效输出") for cid in ("R002","R003","R004","R005","R006")]+[criterion("R007",digest(source_path)==EXPECTED_SHA,"输入哈希") ,criterion("R008",False,"缺少有效输出")]
    else:
        try:
            source=load_workbook(source_path,data_only=False,keep_links=True); output=load_workbook(out_path,data_only=False,keep_links=True)
            structure=output.sheetnames==["LBO","Financials","Operating Cases"]
            results.append(criterion("R001",structure,f"工作表={output.sheetnames}"))
            expected=targets(); target_ok=all(norm(output[s][c].value)==norm(v) for (s,c),v in expected.items())
            results.append(criterion("R002",all(norm(output[s][c].value)==norm(v) for (s,c),v in expected.items() if s=="LBO"),"LBO 目标公式 24 个"))
            results.append(criterion("R003",all(norm(output[s][c].value)==norm(v) for (s,c),v in expected.items() if s=="Financials"),"Financials 目标公式 5 个"))
            regression=True
            for sheet_name,(min_row,max_row,min_col,max_col) in RANGES.items():
                if output[sheet_name].max_row!=source[sheet_name].max_row or output[sheet_name].max_column!=source[sheet_name].max_column: regression=False
                for row in range(min_row,max_row+1):
                    for col in range(min_col,max_col+1):
                        if (sheet_name,output[sheet_name].cell(row,col).coordinate) not in expected and not equal(source[sheet_name].cell(row,col).value,output[sheet_name].cell(row,col).value): regression=False
            results.append(criterion("R004",regression,"非目标答案区域逐格保持"))
            formula_count=sum(isinstance(cell.value,str) and cell.value.startswith("=") for sheet in output.worksheets for row in sheet.iter_rows() for cell in row)
            broken=any(isinstance(cell.value,str) and "#REF!" in cell.value.upper() for sheet in output.worksheets for row in sheet.iter_rows() for cell in row)
            calc=output.calculation
            formula_ok=target_ok and formula_count>=900 and not broken and calc.calcMode=="auto" and bool(calc.fullCalcOnLoad) and bool(calc.forceFullCalc)
            results.append(criterion("R005",formula_ok,f"公式数={formula_count}，#REF!={broken}，自动重算={calc.calcMode}/{calc.fullCalcOnLoad}/{calc.forceFullCalc}"))
            native=structure and formula_count>=900 and all(output[s].protection.sheet is not True for s in output.sheetnames)
            results.append(criterion("R006",native,"工作簿原生可编辑"))
            results.append(criterion("R007",digest(source_path)==EXPECTED_SHA,"输入 SHA256="+digest(source_path)))
            results.append(criterion("R008",digest(out_path)!=digest(source_path) and target_ok,"非未修改输入且目标已修复"))
            source.close(); output.close()
        except Exception as exc:
            results=[criterion("R001",False,f"读取失败：{exc}")]+[criterion(cid,False,"读取失败") for cid in ("R002","R003","R004","R005","R006","R007","R008")]
    values={item["criterion_id"]:item["score"] for item in results}; reward=sum(WEIGHTS[k]*values.get(k,0) for k in WEIGHTS); passed=all(values.get(k,0)==1 for k in ("R001","R002","R003","R004","R005")) and reward>=.8
    payload={"task_id":TASK_ID,"rubric_version":RUBRIC_VERSION,"judge_version":JUDGE_VERSION,"artifact_sha256":[digest(out_path)] if out_path.is_file() else [],"criterion_results":results,"reward":reward,"passed":passed}
    args.result.parent.mkdir(parents=True,exist_ok=True); args.result.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8"); args.reward_file.parent.mkdir(parents=True,exist_ok=True); args.reward_file.write_text(f"{reward:.6f}\n"); raise SystemExit(0 if passed else 1)
if __name__=="__main__": main()
