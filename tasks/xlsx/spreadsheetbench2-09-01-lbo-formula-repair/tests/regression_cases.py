#!/usr/bin/env python3
"""Golden、fresh replay 与七个针对性负例。"""

from __future__ import annotations

import importlib.util
import json
import shutil
import subprocess
import sys
import tempfile
from copy import copy
from pathlib import Path

from openpyxl import load_workbook


TASK_ROOT = Path(__file__).resolve().parents[1]
SOLUTION = TASK_ROOT / "solution"
INPUT = TASK_ROOT / "data" / "input_files" / "09_01_input.xlsx"
OUTPUT_NAMES = ["repair_workbook.py", "09_01_repaired.xlsx", "repair_log.json"]

spec = importlib.util.spec_from_file_location("ssb2_repair_judge", Path(__file__).with_name("judge.py"))
if spec is None or spec.loader is None:
    raise RuntimeError("无法加载 Judge")
judge = importlib.util.module_from_spec(spec)
spec.loader.exec_module(judge)


def copy_candidate(target: Path) -> None:
    target.mkdir(parents=True, exist_ok=True)
    for name in OUTPUT_NAMES:
        shutil.copy2(SOLUTION / name, target / name)


def criterion_scores(result: dict) -> dict[str, float]:
    return {item["criterion_id"]: item["score"] for item in result["criterion_results"]}


def assert_full(label: str, path: Path) -> None:
    result = judge.evaluate(path, render=False)
    assert result["passed"] and result["reward"] == 1.0, (
        label,
        result["reward"],
        result["hurdle_failures"],
    )


def assert_rejected(label: str, path: Path, criterion_id: str) -> None:
    result = judge.evaluate(path, render=False)
    scores = criterion_scores(result)
    assert not result["passed"] and result["reward"] < 0.8, (label, result["reward"])
    assert scores[criterion_id] == 0.0, (label, criterion_id, scores[criterion_id])


def case_golden(base: Path) -> None:
    candidate = base / "golden"
    copy_candidate(candidate)
    assert_full("golden", candidate)


def case_fresh(base: Path) -> None:
    candidate = base / "fresh"
    subprocess.run(
        [
            sys.executable,
            str(SOLUTION / "solve.py"),
            "--input-dir",
            str(INPUT.parent),
            "--output-dir",
            str(candidate),
        ],
        check=True,
        timeout=90,
    )
    assert_full("fresh", candidate)


def case_missing(base: Path) -> None:
    candidate = base / "missing"
    copy_candidate(candidate)
    (candidate / "09_01_repaired.xlsx").unlink()
    assert_rejected("missing workbook", candidate, "R001")


def case_unchanged_input(base: Path) -> None:
    candidate = base / "unchanged"
    copy_candidate(candidate)
    shutil.copy2(INPUT, candidate / "09_01_repaired.xlsx")
    assert_rejected("unchanged input", candidate, "R013")


def case_wrong_target_formula(base: Path) -> None:
    candidate = base / "wrong-target"
    copy_candidate(candidate)
    path = candidate / "09_01_repaired.xlsx"
    workbook = load_workbook(path)
    workbook["LBO"]["G66"] = "=G57+G75"
    workbook.save(path)
    workbook.close()
    assert_rejected("wrong target formula", candidate, "R002")


def case_collateral_regression(base: Path) -> None:
    candidate = base / "collateral"
    copy_candidate(candidate)
    path = candidate / "09_01_repaired.xlsx"
    workbook = load_workbook(path)
    workbook["Financials"]["G17"] = 999
    workbook.save(path)
    workbook.close()
    assert_rejected("collateral value change", candidate, "R003")


def case_hardcoded_processor(base: Path) -> None:
    candidate = base / "hardcoded"
    copy_candidate(candidate)
    script = candidate / "repair_workbook.py"
    script.write_text(
        """#!/usr/bin/env python3
import argparse
import shutil
from pathlib import Path

parser = argparse.ArgumentParser()
parser.add_argument('--input', required=True)
parser.add_argument('--output', required=True, type=Path)
parser.add_argument('--log', required=True, type=Path)
args = parser.parse_args()
source = Path(__file__).resolve().parent
args.output.parent.mkdir(parents=True, exist_ok=True)
args.log.parent.mkdir(parents=True, exist_ok=True)
shutil.copy2(source / '09_01_repaired.xlsx', args.output)
shutil.copy2(source / 'repair_log.json', args.log)
""",
        encoding="utf-8",
    )
    assert_rejected("hardcoded processor", candidate, "R015")


def case_fake_log(base: Path) -> None:
    candidate = base / "fake-log"
    copy_candidate(candidate)
    path = candidate / "repair_log.json"
    payload = json.loads(path.read_text(encoding="utf-8"))
    payload["repaired_formula_cells"] = 3
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    assert_rejected("fake log", candidate, "R008")


def case_style_damage(base: Path) -> None:
    candidate = base / "style-damage"
    copy_candidate(candidate)
    path = candidate / "09_01_repaired.xlsx"
    workbook = load_workbook(path)
    cell = workbook["Financials"]["B10"]
    font = copy(cell.font)
    font.bold = not bool(font.bold)
    cell.font = font
    workbook.save(path)
    workbook.close()
    assert_rejected("style damage", candidate, "R006")


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="cowork-evo-ssb2-regression-") as directory:
        base = Path(directory)
        case_golden(base)
        case_fresh(base)
        case_missing(base)
        case_unchanged_input(base)
        case_wrong_target_formula(base)
        case_collateral_regression(base)
        case_hardcoded_processor(base)
        case_fake_log(base)
        case_style_damage(base)
    print("PASS: Golden=1.0, fresh replay=1.0, 7/7 negative cases rejected")


if __name__ == "__main__":
    main()
