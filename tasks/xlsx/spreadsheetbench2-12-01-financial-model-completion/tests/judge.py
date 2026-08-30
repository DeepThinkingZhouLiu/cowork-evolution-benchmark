#!/usr/bin/env python3
from __future__ import annotations
import argparse, hashlib, json, shutil, subprocess, tempfile, zipfile
from pathlib import Path
from openpyxl import load_workbook

OUT="financial_model_completed.xlsx"; TASK_ID="cowork-evo/spreadsheetbench2-12-01-financial-model-completion"
WEIGHTS={"R001":.15,"R002":.30,"R003":.20,"R004":.20,"R005":.10,"R006":.05}
def crit(cid,ok,ev): return {"criterion_id":cid,"score":1.0 if ok else 0.0,"evidence":ev,"method":"deterministic"}
def recalc(path: Path):
    lo=shutil.which("libreoffice") or shutil.which("soffice")
    if not lo: return None
    with tempfile.TemporaryDirectory(prefix="sbench-12-01-") as td:
        r=subprocess.run([lo,"--headless","--convert-to","xlsx","--outdir",td,str(path)],capture_output=True,text=True,timeout=180,check=False)
        p=Path(td)/path.name
        if r.returncode!=0 or not p.is_file(): return None
        data=p.read_bytes(); return load_workbook(__import__('io').BytesIO(data),data_only=True)
def main():
    ap=argparse.ArgumentParser(); ap.add_argument("--artifact-dir",type=Path,required=True); ap.add_argument("--result",type=Path,required=True); ap.add_argument("--reward-file",type=Path,required=True); a=ap.parse_args(); source=Path(__file__).parents[1]/"data/input_files/financial_model_input.xlsx"; out=a.artifact_dir/OUT; results=[]
    try:
        native=out.is_file() and zipfile.is_zipfile(out); results.append(crit("R001",native,"输出为可打开的 OOXML 工作簿"))
        wi=load_workbook(source,data_only=False); wo=load_workbook(out,data_only=False) if native else None
        expected={}
        for col,prev in zip("IJKLM","HIJKL"): expected[("Assumption",f"{col}8")]=f"={prev}8*(1+{col}9)"; expected[("Assumption",f"{col}18")]=f"={col}14*{col}19"; expected[("Working Capital",f"{col}6")]=f"={col}7*{col}8/365"; expected[("Working Capital",f"{col}8")]=70 if col=="I" else f"={prev}8"
        for col,prev in zip("IJKLMN","HIJKLM"): expected[("Working Capital",f"{col}30")]=f"={prev}29-{col}29"
        for col in "CDEFGHIJKLM": expected[("P&L",f"{col}37")]=f"={col}36/{col}9"
        formulas_ok=bool(wo) and all(wo[s][c].value==v for (s,c),v in expected.items()); results.append(crit("R002",formulas_ok,f"目标公式数={len(expected)}"))
        changed=set(expected)
        preserved=True
        if wo:
            for s in wi.sheetnames:
                if s not in wo.sheetnames: preserved=False; break
                w1,w2=wi[s],wo[s]
                if (w1.max_row,w1.max_column)!=(w2.max_row,w2.max_column): preserved=False; break
                for row in range(1,w1.max_row+1):
                    for col in range(1,w1.max_column+1):
                        if (s,w1.cell(row,col).coordinate) in changed: continue
                        left, right = w1.cell(row,col).value, w2.cell(row,col).value
                        if isinstance(left, (int, float)) and isinstance(right, (int, float)):
                            same = abs(float(left) - float(right)) <= 1e-9
                        else:
                            same = left == right
                        if not same: preserved=False; break
                    if not preserved: break
                if not preserved: break
        else: preserved=False
        results.append(crit("R003",preserved,"非目标单元格逐格保留"))
        calc=recalc(out) if native else None; calc_ok=False
        if calc:
            calc_ok=all(calc[s][c].value is not None and not (isinstance(calc[s][c].value,str) and calc[s][c].value.startswith("#")) for s,c in expected)
            calc.close()
        results.append(crit("R004",calc_ok,"LibreOffice 重算后目标单元格均为有效数值"))
        results.append(crit("R005",bool(wo and wo.calculation.calcMode=="auto" and wo.calculation.fullCalcOnLoad),"工作簿启用自动重算"))
        results.append(crit("R006",hashlib.sha256(source.read_bytes()).hexdigest()=="7f258ae78f7b8b7856a56b91a3a30632d1b11f02e4a50dd59be2d80e89fb42fd","输入 SHA256 固定"))
        wi.close();
        if wo: wo.close()
    except Exception as exc: results=[crit("R001",False,f"Judge 读取失败：{exc}")]
    vals={x["criterion_id"]:x["score"] for x in results}; reward=sum(WEIGHTS[k]*vals.get(k,0) for k in WEIGHTS); passed=reward>=.8 and vals.get("R001",0)==1 and vals.get("R002",0)==1 and vals.get("R003",0)==1
    payload={"task_id":TASK_ID,"rubric_version":"spreadsheetbench2-12-01-v0.1","judge_version":"spreadsheetbench2-12-01-deterministic-v0.1","artifact_sha256":[hashlib.sha256(out.read_bytes()).hexdigest()] if out.is_file() else [],"criterion_results":results,"reward":reward,"passed":passed}; a.result.parent.mkdir(parents=True,exist_ok=True); a.result.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8"); a.reward_file.parent.mkdir(parents=True,exist_ok=True); a.reward_file.write_text(f"{reward:.6f}\n"); raise SystemExit(0 if passed else 1)
if __name__=="__main__": main()
