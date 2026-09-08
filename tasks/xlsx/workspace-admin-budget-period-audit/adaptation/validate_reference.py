"""预算题专属验收：窗口、重复映射、单位、空年度、公式依赖反例。"""
import argparse,json,shutil,subprocess,sys
from pathlib import Path
from openpyxl import Workbook,load_workbook
def main():
    p=argparse.ArgumentParser();p.add_argument('--output',type=Path,required=True);a=p.parse_args();out=a.output.resolve()
    out.mkdir(parents=True,exist_ok=False);root=Path(__file__).resolve().parents[1]
    fresh=out/'fresh';subprocess.run([sys.executable,str(root/'solution/solve.py'),'--input-dir',str(root/'data/input_files'),'--output-dir',str(fresh)],check=True)
    cases={'fresh':fresh,'golden':root/'solution/Golden'}
    for label in ['empty','blank','corrupt','unchanged','unit_error','duplicate_budget','wrong_window','lost_blank','constant_formulas']:
        folder=out/label;folder.mkdir();cases[label]=folder;f=folder/'admin_audit.xlsx'
        if label=='empty':continue
        if label=='blank':Workbook().save(f);continue
        if label=='corrupt':f.write_bytes(b'broken');continue
        if label=='unchanged':shutil.copy2(root/'data/input_files/1-行政预算及支出管理表.xlsx',f);continue
        w=load_workbook(fresh/'admin_audit.xlsx')
        if label=='unit_error':
            for r in range(2,37):
                w['预算2025'][f'P{r}']=f'=SUM(D{r}:O{r})';w['预算2025'][f'Q{r}']=f'=G{r}'
        if label=='duplicate_budget':w['审计指标']['B22']="=SUM('预算对比'!G2:G13)"
        if label=='wrong_window':w['审计指标']['B15']='=54168.8'
        if label=='lost_blank':w['预算2025']['C3']=7.2
        if label=='constant_formulas':
            cached=out/'cached';cached.mkdir()
            subprocess.run(['libreoffice',f'-env:UserInstallation={(out/"cache-profile").as_uri()}','--headless','--convert-to','xlsx','--outdir',str(cached),str(fresh/'admin_audit.xlsx')],check=True,capture_output=True)
            values=load_workbook(cached/'admin_audit.xlsx',data_only=True)
            for s in w:
                for row in s:
                    for c in row:
                        if c.data_type=='f':
                            value=values[s.title][c.coordinate].value;c.value='='+('"'+value+'"' if isinstance(value,str) else str(value))
        w.save(f)
    scores={}
    for label,folder in cases.items():
        logdir=out/'judges'/label;logdir.mkdir(parents=True)
        with (logdir/'judge.log').open('w') as log:
            run=subprocess.run([sys.executable,str(root/'tests/judge.py'),'--artifact-dir',str(folder),'--result',str(logdir/'result.json'),'--reward-file',str(logdir/'reward.txt')],stdout=log,stderr=subprocess.STDOUT,timeout=300)
        assert run.returncode in (0,1),(label,run.returncode)
        scores[label]=json.loads((logdir/'result.json').read_text())['reward'];print(label,scores[label],flush=True)
    assert scores['fresh']==scores['golden']==1
    assert all(scores[x]==0 for x in ['empty','blank','corrupt','unchanged'])
    assert all(0<scores[x]<1 for x in ['unit_error','duplicate_budget','wrong_window','lost_blank','constant_formulas'])
    assert scores['constant_formulas']<.8 and scores['unit_error']<.8
    repeats=[]
    for n in range(3):
        folder=out/f'repeat-{n}';folder.mkdir()
        subprocess.run([sys.executable,str(root/'tests/judge.py'),'--artifact-dir',str(fresh),'--result',str(folder/'result.json'),'--reward-file',str(folder/'reward.txt')],check=True,capture_output=True,timeout=300)
        repeats.append(json.loads((folder/'result.json').read_text())['reward'])
    assert repeats==[1]*3
    (out/'validation.json').write_text(json.dumps({'task':'workspace-admin-budget-period-audit','scores':scores,'repeated_rewards':repeats,'all_assertions_passed':True},indent=2,ensure_ascii=False))
if __name__=='__main__':main()
