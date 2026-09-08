"""行政预算题的独立参考解；只处理已逐份审查的四份预算/备用金材料。"""
import argparse
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.chart import BarChart,Reference
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--input-dir',type=Path,required=True);parser.add_argument('--output-dir',type=Path,required=True);a=parser.parse_args()
    w=Workbook();w.remove(w.active)
    s=w.create_sheet('实际2024');s.append(['日期','原项目','金额元','部门','来源'])
    source=load_workbook(a.input_dir/'1-行政预算及支出管理表.xlsx',data_only=True).active
    for r in range(20,33):s.append([source.cell(r,c).value for c in (4,5,8,9)]+[f'1-行政预算及支出管理表.xlsx!{source.title}!D{r}:I{r}'])
    s=w.create_sheet('预算2025');s.append(['大类','子项','原年度万元']+[f'{m}月万元' for m in range(1,13)]+['推导年度元','四月元','来源'])
    budget=load_workbook(a.input_dir/'2025年度行政费用预算表.xlsx',data_only=True).active;group=''
    for r in range(4,46):
        if budget.cell(r,1).value:group=budget.cell(r,1).value
        if not budget.cell(r,2).value:continue
        n=s.max_row+1
        s.append([group,budget.cell(r,2).value,budget.cell(r,4).value]+[budget.cell(r,c).value or 0 for c in range(5,17)]+[f'=SUM(D{n}:O{n})*10000',f'=G{n}*10000',f'2025年度行政费用预算表.xlsx!{budget.title}!D{r}:P{r}'])
    s=w.create_sheet('备用2022');s.append(['日期','部门','经手人','事项','收支','金额元','报销日期','来源'])
    cash=load_workbook(a.input_dir/'备用金收支明细表.xlsx',data_only=True).active
    for r in range(5,10):
        reimbursement=cash.cell(r,16).value
        s.append([cash.cell(r,c).value for c in (9,10,11,12,13,14)]+[reimbursement if isinstance(reimbursement,datetime) else None,f'备用金收支明细表.xlsx!{cash.title}!I{r}:Q{r}'])
    s=w.create_sheet('备用2024');s.append(['日期','类型','金额元','事项','来源'])
    cash=load_workbook(a.input_dir/'备用金收支结余明细表.xlsx',data_only=True).active
    for kind,cols in [('存入',(2,5,3)),('支出',(6,8,7))]:
        for r in range(8,24):s.append([cash.cell(r,cols[0]).value,kind,cash.cell(r,cols[1]).value,cash.cell(r,cols[2]).value,f'备用金收支结余明细表.xlsx!{cash.title}!{get_column_letter(cols[0])}{r}:{get_column_letter(cols[1])}{r}'])
    s=w.create_sheet('预算对比');s.append(['原项目','原预算元','四月实际元','实际减原预算','超支','情景映射','2025四月预算元','实际减情景预算'])
    mapping={'办公费':'办公用品','交通费':'差旅费','招待费':'宴请费','广告费':'宣传费','维修费':'设备修理费','租赁费':'办公室租金','水电费':'水电费','物业费':'物业费','差旅费':'差旅费'}
    for r in range(20,32):
        n=s.max_row+1;item=source.cell(r,11).value;m=mapping.get(item,'未映射')
        s.append([item,source.cell(r,12).value,f'=SUMIF(\'实际2024\'.B$2:B$14,A{n},\'实际2024\'.C$2:C$14)',f'=C{n}-B{n}',f'=IF(D{n}>0,"是","否")',m,f'=SUMIF(\'预算2025\'.B$2:B$36,F{n},\'预算2025\'.Q$2:Q$36)' if m!='未映射' else None,f'=C{n}-G{n}' if m!='未映射' else None])
    s=w.create_sheet('审计指标');s.append(['指标','结果元/项目数'])
    # 日期条件必须分别作用于左右两组交易，不能使用源表已混口径的余额。
    def cash_sum(kind,inside=True):
        base=f'SUMIFS(\'备用2024\'.C$2:C$33,\'备用2024\'.B$2:B$33,"{kind}",\'备用2024\'.A$2:A$33,">="&DATE(2024,5,1),\'备用2024\'.A$2:A$33,"<="&DATE(2024,5,15))'
        return base if inside else f'SUMIF(\'备用2024\'.B$2:B$33,"{kind}",\'备用2024\'.C$2:C$33)-{base}'
    metrics=[
        ('2024四月实际合计',"SUM('实际2024'.C2:C14)"),('原预算合计',"SUM('预算对比'.B2:B13)"),('原预算减实际','B3-B2'),('超支项目数','COUNTIF(\'预算对比\'.E2:E13,"是")'),
        ('2025年度预算元',"SUM('预算2025'.P2:P36)"),('2025四月预算元',"SUM('预算2025'.Q2:Q36)"),('2022期初备用金','50000'),
        ('2022支出绝对额','-SUMIF(\'备用2022\'.E2:E6,"支",\'备用2022\'.F2:F6)'),('2022存入','SUMIF(\'备用2022\'.E2:E6,"收",\'备用2022\'.F2:F6)'),('2022期末','B8-B9+B10'),
        ('2022尚未报销支出','-SUMIFS(\'备用2022\'.F2:F6,\'备用2022\'.E2:E6,"支",\'备用2022\'.G2:G6,"")'),
        ('2024窗口存入',cash_sum('存入')),('2024窗口支出',cash_sum('支出')),('2024窗口净额','B13-B14'),('源表展示余额','54168.8'),('展示余额减窗口净额','B16-B15'),
        ('窗口外存入',cash_sum('存入',False)),('窗口外支出',cash_sum('支出',False)),('窗口外净额','B18-B19'),
        ('已映射实际','SUMIF(\'预算对比\'.F2:F13,"<>未映射",\'预算对比\'.C2:C13)'),
        ('去重后情景预算','SUM(\'预算对比\'.G2:G13)-SUMIF(\'预算对比\'.A2:A13,"交通费",\'预算对比\'.G2:G13)'),
        ('未映射实际','SUMIF(\'预算对比\'.F2:F13,"未映射",\'预算对比\'.C2:C13)'),
        ('已映射实际减去重预算','B21-B22'),('共同差旅实际','SUMIF(\'预算对比\'.F2:F13,"差旅费",\'预算对比\'.C2:C13)'),
        ('共同差旅预算','SUMIF(\'预算2025\'.B2:B36,"差旅费",\'预算2025\'.Q2:Q36)')]
    for label,f in metrics:s.append([label,'='+f.replace("'.","'!")])
    # openpyxl 写入 OOXML 标准引用，避免依赖某款表格软件的输入语法。
    for sheet in w:
        for row in sheet:
            for c in row:
                if c.data_type=='f':c.value=c.value.replace("'.","'!")
    s=w.create_sheet('问题清单');s.append(['主题','原证据','数值影响','处理与待核实'])
    for row in [
        ['跨期','行政支出2024-04；备用金2022-01、2024-05；预算2025','四套期间不重叠','预算对比仅为情景，不宣称2025全年剩余预算，需另取同年实际'],
        ['单位与空年度','2025源表单位万元，物业费/保洁费年度原格空白','按十二个月补算；保留原空白不冒充源值','叶子行求和；排除大类小计，不以年度/12替代季节性四月'],
        ['窗口余额','备用表I3全范围；E5/I5为5月1日至15日','54168.8-54129.8=39；区间外净额39','分别按存入/支出日期重算，净额不包含未知期初'],
        ['共享预算','交通费与差旅费按题设均映射差旅费','合并实际；情景预算只取一次','映射为分析假设，不能把两行预算相加当可用额度'],
        ['不可映射项目','福利费、保险费、会议费没有题设可靠对应','未映射实际单列5300元','保留待映射，不编造审批，不将保险随意映射车险']]:s.append(row)
    chart=BarChart();chart.title='2024四月原预算与实际（元）';chart.add_data(Reference(w['预算对比'],min_col=2,max_col=3,min_row=1,max_row=13),titles_from_data=True);chart.set_categories(Reference(w['预算对比'],min_col=1,min_row=2,max_row=13));chart.width=24;chart.height=11;w['预算对比'].add_chart(chart,'J2')
    for s in w:
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
        for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='24476B')
        for row in s.iter_rows(min_row=2):
            for c in row:
                c.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(c.value,datetime):c.number_format='yyyy-mm-dd'
                elif isinstance(c.value,(int,float)) or c.data_type=='f':c.number_format='#,##0.00'
        for col in range(1,s.max_column+1):s.column_dimensions[get_column_letter(col)].width=22 if col<4 else 18
        if s.title=='问题清单':
            for col in 'BCD':s.column_dimensions[col].width=48
            for i in range(2,7):s.row_dimensions[i].height=60
        if s.title in ('实际2024','备用2022','备用2024'):
            source_col={'实际2024':'E','备用2022':'H','备用2024':'E'}[s.title]
            s.column_dimensions[source_col].width=60
            for r in range(2,s.max_row+1):s.row_dimensions[r].height=42
        if s.title=='预算2025':
            s.column_dimensions['R'].width=48
            for c in range(4,16):s.column_dimensions[get_column_letter(c)].width=10
            for r in range(2,s.max_row+1):s.row_dimensions[r].height=42
        if s.title=='审计指标':s.column_dimensions['A'].width=34
        s.sheet_properties.pageSetUpPr.fitToPage=True;s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0;s.print_title_rows='1:1'
    a.output_dir.mkdir(parents=True,exist_ok=True);w.save(a.output_dir/'admin_audit.xlsx')
if __name__=='__main__':main()
