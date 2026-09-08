"""预算题专属阅卷：独立读取原四文件，核对期间、单位、共享预算并扰动复算。"""
import argparse,json,math,subprocess,tempfile,zipfile
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

def near(x,y):
    return isinstance(x,(int,float)) and not isinstance(x,bool) and math.isclose(x,y,rel_tol=1e-8,abs_tol=.005)
def recalc(path,root,name):
    dest=root/name;dest.mkdir();profile=root/(name+'-profile')
    run=subprocess.run(['libreoffice',f'-env:UserInstallation={profile.as_uri()}','--headless','--convert-to','xlsx','--outdir',str(dest),str(path)],capture_output=True,timeout=100)
    if run.returncode or not (dest/path.name).exists():raise RuntimeError('LibreOffice未完成重算')
    return load_workbook(dest/path.name,data_only=True)
def main():
    p=argparse.ArgumentParser();p.add_argument('--artifact-dir',type=Path,required=True);p.add_argument('--result',type=Path,required=True);p.add_argument('--reward-file',type=Path,required=True);a=p.parse_args()
    tests=Path(__file__).parent;inp=tests.parent/'data/input_files';rub=json.loads((tests/'rubric.json').read_text());path=a.artifact_dir/'admin_audit.xlsx'
    scores={c['id']:0. for c in rub['criteria']};evidence={}
    def mark(k,checks,why):scores[k]=sum(bool(x) for x in checks)/len(checks) if checks else 0.;evidence[k]=why
    w=None
    try:
        if path.is_file():w=load_workbook(path)
    except (OSError,ValueError,KeyError,zipfile.BadZipFile):pass
    if w:
        mark('R001',['实际2024' in w and w['实际2024'].max_row>1],'原生有效工作簿及非空业务底稿')
        actual=load_workbook(inp/'1-行政预算及支出管理表.xlsx',data_only=True).active
        rows=[[actual.cell(r,c).value for c in (4,5,8,9)] for r in range(20,33)]
        def raw(sheet,expected,src_col):
            if sheet not in w:return [False]
            s=w[sheet];checks=[s.max_row==len(expected)+1]
            for i,row in enumerate(expected,2):
                for j,value in enumerate(row,1):
                    got=s.cell(i,j).value
                    checks.append(near(got,value) if isinstance(value,(int,float)) else got==value)
                loc=str(s.cell(i,src_col).value or '');checks.append('.xlsx!' in loc)
            return checks
        mark('R002',raw('实际2024',rows,5),'13条原始实际、日期与来源；不计模板空行')
        budget=load_workbook(inp/'2025年度行政费用预算表.xlsx',data_only=True).active
        leaves=[];group=''
        for r in range(4,46):
            if budget.cell(r,1).value:group=budget.cell(r,1).value
            if budget.cell(r,2).value:leaves.append([group,budget.cell(r,2).value,budget.cell(r,4).value]+[budget.cell(r,c).value or 0 for c in range(5,17)])
        mark('R003',raw('预算2025',leaves,18),'35叶子科目、原空白、12个月万元；保留大类避免其他科目混淆')
        old=load_workbook(inp/'备用金收支明细表.xlsx',data_only=True).active
        cash22=[]
        for r in range(5,10):cash22.append([old.cell(r,c).value for c in (9,10,11,12,13,14)]+[old.cell(r,16).value if isinstance(old.cell(r,16).value,datetime) else None])
        new=load_workbook(inp/'备用金收支结余明细表.xlsx',data_only=True).active
        cash24=[]
        for kind,cols in [('存入',(2,5,3)),('支出',(6,8,7))]:
            for r in range(8,24):cash24.append([new.cell(r,cols[0]).value,kind,new.cell(r,cols[1]).value,new.cell(r,cols[2]).value])
        mark('R004',raw('备用2022',cash22,8)+raw('备用2024',cash24,5),'分开两个年份，32条左右账记录保留区间外，不错误裁剪或换符号')
        mapping={'办公费':'办公用品','交通费':'差旅费','招待费':'宴请费','广告费':'宣传费','维修费':'设备修理费','租赁费':'办公室租金','水电费':'水电费','物业费':'物业费','差旅费':'差旅费'}
        originals={actual.cell(r,11).value:actual.cell(r,12).value for r in range(20,32)}
        costs={k:sum(x[2] for x in rows if x[1]==k) for k in originals}
        apr={x[1]:x[6]*10000 for x in leaves if x[1]!='其他'}
        def csum(kind,inside):
            return sum(x[2] for x in cash24 if x[1]==kind and (datetime(2024,5,1)<=x[0]<=datetime(2024,5,15))==inside)
        expected=[sum(costs.values()),sum(originals.values()),sum(originals.values())-sum(costs.values()),sum(costs[k]>originals[k] for k in originals),
                  sum(sum(x[3:15])*10000 for x in leaves),sum(x[6]*10000 for x in leaves),old['D5'].value,
                  -sum(x[5] for x in cash22 if x[4]=='支'),sum(x[5] for x in cash22 if x[4]=='收'),old['D5'].value+sum(x[5] for x in cash22),
                  -sum(x[5] for x in cash22 if x[4]=='支' and x[6] is None),csum('存入',True),csum('支出',True),csum('存入',True)-csum('支出',True),
                  new['I3'].value,new['I3'].value-csum('存入',True)+csum('支出',True),csum('存入',False),csum('支出',False),csum('存入',False)-csum('支出',False),
                  sum(costs[k] for k in mapping),sum(apr[k] for k in set(mapping.values())),sum(costs[k] for k in costs if k not in mapping),
                  sum(costs[k] for k in mapping)-sum(apr[k] for k in set(mapping.values())),costs['交通费']+costs['差旅费'],apr['差旅费']]
        required=['实际2024','预算2025','备用2022','备用2024','预算对比','审计指标','问题清单']
        with tempfile.TemporaryDirectory(prefix='admin-budget-judge-') as td:
            temp=Path(td)
            try:v=recalc(path,temp,'baseline')
            except (RuntimeError,subprocess.TimeoutExpired) as err:
                a.result.parent.mkdir(parents=True,exist_ok=True);a.result.write_text(json.dumps({'status':'infrastructure_failure','error':str(err)},ensure_ascii=False));raise SystemExit(2)
            checks=[]
            if '预算对比' in v:
                s=v['预算对比'];checks.append(s.max_row==13)
                for i,k in enumerate(originals,2):
                    want=[k,originals[k],costs[k],costs[k]-originals[k],'是' if costs[k]>originals[k] else '否',mapping.get(k,'未映射'),apr[mapping[k]] if k in mapping else None,costs[k]-apr[mapping[k]] if k in mapping else None]
                    for j,val in enumerate(want,1):
                        got=s.cell(i,j).value;checks.append(near(got,val) if isinstance(val,(int,float)) else got==val)
            mark('R005',checks,'12类原预算/实际/映射/情景；未映射须空白而非假数')
            checks=[near(v['审计指标'].cell(i,2).value,x) if '审计指标' in v else False for i,x in enumerate(expected,2)]
            if '预算2025' in v:
                for i,row in enumerate(leaves,2):checks.extend([near(v['预算2025'].cell(i,16).value,sum(row[3:15])*10000),near(v['预算2025'].cell(i,17).value,row[6]*10000)])
            mark('R006',checks,'25个指标与35行单位转换实际重算；预算去重和日期窗口分别核对')
            coords=[('预算2025',f'{c}{r}') for r in range(2,37) for c in 'PQ']+[('审计指标',f'B{r}') for r in range(2,27)]+[('预算对比',f'{c}{i}') for i,k in enumerate(originals,2) for c in ('CDEGH' if k in mapping else 'CDE')]
            mark('R007',[s in w and w[s][c].data_type=='f' for s,c in coords],'原生公式，不限制等价实现')
            probe=[]
            if all(s in w for s in required):
                q=load_workbook(path)
                if near(q['实际2024']['C2'].value,1100) and near(q['预算2025']['G2'].value,5):
                    q['实际2024']['C2']=1200;q['预算2025']['G2']=5.25
                    q['备用2024']['C17']=1500
                    # 修改月度预算/实际/窗口外金额，然后重排两个底稿的首尾记录，检验按项目关联。
                    for s,n in [(q['实际2024'],14),(q['预算2025'],36)]:
                        left=[s.cell(2,c).value for c in range(1,s.max_column+1)]
                        right=[s.cell(n,c).value for c in range(1,s.max_column+1)]
                        # 预算行有相对公式，翻译公式后移动；不破坏候选自身正确的相对引用。
                        from openpyxl.formula.translate import Translator
                        for c in range(1,s.max_column+1):
                            lv=left[c-1];rv=right[c-1]
                            if isinstance(rv,str) and rv.startswith('='):rv=Translator(rv,origin=s.cell(n,c).coordinate).translate_formula(s.cell(2,c).coordinate)
                            if isinstance(lv,str) and lv.startswith('='):lv=Translator(lv,origin=s.cell(2,c).coordinate).translate_formula(s.cell(n,c).coordinate)
                            s.cell(2,c,rv);s.cell(n,c,lv)
                    probe_path=temp/'changed.xlsx';q.save(probe_path)
                    try:pv=recalc(probe_path,temp,'changed')
                    except (RuntimeError,subprocess.TimeoutExpired) as err:
                        a.result.parent.mkdir(parents=True,exist_ok=True);a.result.write_text(json.dumps({'status':'infrastructure_failure','error':str(err)},ensure_ascii=False));raise SystemExit(2)
                    desired={2:expected[0]+100,6:expected[4]+2500,7:expected[5]+2500,13:expected[11],15:expected[13],18:1500,20:1500-expected[17],21:expected[19]+100,22:expected[20]+2500,24:expected[22]-2400}
                    stable_ok=all(near(pv['审计指标'].cell(r,2).value,desired[r]) for r in (13,15))
                    probe=[near(pv['审计指标'].cell(r,2).value,val) and stable_ok for r,val in desired.items() if r not in (13,15)]
                    byname={r[0]:r for r in pv['预算对比'].iter_rows(min_row=2,max_col=8,values_only=True)}
                    probe.extend([near(byname.get('办公费',[None]*8)[2],1320),near(byname.get('租赁费',[None]*8)[6],52500)])
            mark('R008',probe,'实际、预算、窗口外收款变更及底稿重排后，指标/关联更新；未运行则不得分')
        lines=[' '.join(str(x or '') for x in row) for row in w['问题清单'].iter_rows(min_row=2,max_col=4,values_only=True)] if '问题清单' in w else []
        groups=[('跨期',['2022','2024','2025']),('单位',['万元','空白']),('窗口',['39','期初']),('共享',['交通','差旅','一次']),('映射',['福利','保险','会议'])]
        mark('R009',[any(topic in row and all(word in row for word in words) and len(row)>35 for row in lines) for topic,words in groups],'五类证据与风险覆盖；只判客观要点，不声称语义/审美全评')
        mark('R010',[any(s._charts for s in w),not w._external_links,all(s in w and w[s].freeze_panes and w[s].auto_filter.ref for s in required)],'图表、无外链和筛选冻结')
    reward=sum(c['weight']*scores[c['id']] for c in rub['criteria'])
    if scores['R001']<1:reward=0
    reward=round(max(0,min(1,reward)),6)
    result={'task_id':rub['task_id'],'rubric_version':rub['rubric_version'],'judge_version':'admin-budget-period-v1','reward':reward,'passed':reward>=rub['pass_threshold'],'criterion_results':[{'criterion_id':k,'score':round(v,6),'evidence':evidence.get(k,'未满足'),'method':'deterministic'} for k,v in scores.items()]}
    a.result.parent.mkdir(parents=True,exist_ok=True);a.result.write_text(json.dumps(result,ensure_ascii=False,indent=2));a.reward_file.parent.mkdir(parents=True,exist_ok=True);a.reward_file.write_text(f'{reward:.6f}\n')
    raise SystemExit(0 if result['passed'] else 1)
if __name__=='__main__':main()
