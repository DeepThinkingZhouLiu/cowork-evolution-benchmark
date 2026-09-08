"""年度工资题专属参考与反例测试。"""
import argparse,json,shutil,subprocess,sys
from pathlib import Path
from openpyxl import Workbook,load_workbook
def main():
    p=argparse.ArgumentParser();p.add_argument('--output',type=Path,required=True);a=p.parse_args();out=a.output.resolve();out.mkdir(parents=True,exist_ok=False);root=Path(__file__).resolve().parents[1]
    fresh=out/'fresh';subprocess.run([sys.executable,str(root/'solution/solve.py'),'--input-dir',str(root/'data/input_files'),'--output-dir',str(fresh)],check=True)
    cases={'fresh':fresh,'golden':root/'solution/Golden'}
    for name in ['empty','blank','corrupt','unchanged','keep_double_tax','sample_std','overwrite_annual','lost_month','constant_formulas']:
        folder=out/name;folder.mkdir();cases[name]=folder;f=folder/'annual_payroll_audit.xlsx'
        if name=='empty':continue
        if name=='blank':Workbook().save(f);continue
        if name=='corrupt':f.write_bytes(b'broken');continue
        if name=='unchanged':shutil.copy2(next((root/'data/input_files').glob('*.xlsx')),f);continue
        w=load_workbook(fresh/'annual_payroll_audit.xlsx')
        if name=='keep_double_tax':
            for r in range(2,98):w['月度勾稽'][f'F{r}']=f'=G{r}'
        if name=='sample_std':
            for r in range(2,10):w['员工年度'][f'N{r}']=w['员工年度'][f'N{r}'].value.replace('STDEVP','STDEV')
        if name=='overwrite_annual':w['年度底稿']['Q2']=1
        if name=='lost_month':
            for row in w['月度底稿'].iter_rows(min_row=90,max_row=97,min_col=8,max_col=24):
                for c in row:
                    if isinstance(c.value,(int,float)):c.value=0
        if name=='constant_formulas':
            cached=out/'cached';cached.mkdir()
            subprocess.run(['libreoffice',f'-env:UserInstallation={(out/"cache-profile").as_uri()}','--headless','--convert-to','xlsx','--outdir',str(cached),str(fresh/'annual_payroll_audit.xlsx')],check=True,capture_output=True)
            values=load_workbook(cached/'annual_payroll_audit.xlsx',data_only=True)
            for s in w:
                for row in s:
                    for c in row:
                        if c.data_type=='f':c.value='='+str(values[s.title][c.coordinate].value)
        w.save(f)
    scores={}
    for name,folder in cases.items():
        j=out/'judges'/name;j.mkdir(parents=True)
        with (j/'judge.log').open('w') as log:r=subprocess.run([sys.executable,str(root/'tests/judge.py'),'--artifact-dir',str(folder),'--result',str(j/'result.json'),'--reward-file',str(j/'reward.txt')],stdout=log,stderr=subprocess.STDOUT,timeout=360)
        assert r.returncode in (0,1),(name,r.returncode)
        scores[name]=json.loads((j/'result.json').read_text())['reward'];print(name,scores[name],flush=True)
    assert scores['fresh']==scores['golden']==1
    assert all(scores[x]==0 for x in ['empty','blank','corrupt','unchanged'])
    assert all(0<scores[x]<1 for x in ['keep_double_tax','sample_std','overwrite_annual','lost_month'])
    assert scores['constant_formulas']<.8
    repeated=[]
    for i in range(3):
        j=out/f'repeat-{i}';j.mkdir()
        subprocess.run([sys.executable,str(root/'tests/judge.py'),'--artifact-dir',str(fresh),'--result',str(j/'result.json'),'--reward-file',str(j/'reward.txt')],check=True,capture_output=True,timeout=360);repeated.append(json.loads((j/'result.json').read_text())['reward'])
    assert repeated==[1]*3
    (out/'validation.json').write_text(json.dumps({'task':'workspace-annual-payroll-ledger-audit','scores':scores,'repeated_rewards':repeated,'all_assertions_passed':True},ensure_ascii=False,indent=2))
if __name__=='__main__':main()

