"""年度工资题参考解：96月人记录独立勾稽，不重新估税。"""
import argparse
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.chart import BarChart,Reference
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter as col
MONTHS=['一月','二月','三月','四月','五月','六月','七月','八月','九月','十月','十一月','十二月']
def main():
    p=argparse.ArgumentParser();p.add_argument('--input-dir',type=Path,required=True);p.add_argument('--output-dir',type=Path,required=True);a=p.parse_args()
    filename='完整年度工资管理系统_2019.xlsx';src=load_workbook(a.input_dir/filename,data_only=True);w=Workbook();w.remove(w.active)
    raw=w.create_sheet('月度底稿');raw.append(['月份']+[src['一月'].cell(1,c).value for c in range(1,25)]+['来源'])
    for month,name in enumerate(MONTHS,1):
        for r in range(2,10):raw.append([month]+[src[name].cell(r,c).value for c in range(1,25)]+[f'{filename}!{name}!A{r}:X{r}'])
    annual=w.create_sheet('年度底稿');annual.append([src['年度应发工资汇总'].cell(1,c).value for c in range(1,19)]+['来源'])
    for r in range(2,10):annual.append([src['年度应发工资汇总'].cell(r,c).value for c in range(1,19)]+[f'{filename}!年度应发工资汇总!A{r}:R{r}'])
    audit=w.create_sheet('月度勾稽');audit.append(['月份','序号','姓名','重算应发','重算扣款','单次扣款净额','原实发','情景差额','应发残差','扣款残差','税前残差','实发链残差'])
    def original(letter,r):return f'SUMIFS(\'月度底稿\'!{letter}$2:{letter}$97,\'月度底稿\'!A$2:A$97,A{r},\'月度底稿\'!B$2:B$97,B{r})'
    for row in raw.iter_rows(min_row=2,values_only=True):
        r=audit.max_row+1
        audit.append([row[0],row[1],row[2],'='+ '+'.join(original(c,r) for c in 'HIJKLMN'),'='+ '+'.join(original(c,r) for c in 'PQRST'),f'=D{r}-E{r}','='+original('W',r),f'=F{r}-G{r}','='+original('O',r)+f'-D{r}','='+original('U',r)+f'-E{r}','='+original('V',r)+f'-(D{r}-E{r})','='+original('W',r)+'-('+original('V',r)+'-'+original('T',r)+')'])
    emp=w.create_sheet('员工年度');emp.append(['序号','姓名','部门','原实发合计','源年度Q','勾稽差额','重算应发','单次扣款情景','情景增额','月均实发','最低月实发','最高月实发','极差','总体标准差','变异系数'])
    for source_row in src['基础信息表'].iter_rows(min_row=2,max_col=4,values_only=True):
        ident,name,dept,_=source_row;r=emp.max_row+1
        monthly=[f'SUMIFS(\'月度勾稽\'!G$2:G$97,\'月度勾稽\'!B$2:B$97,A{r},\'月度勾稽\'!A$2:A$97,{month})' for month in range(1,13)]
        lookup=lambda c:f'SUMIF(\'月度勾稽\'!B$2:B$97,A{r},\'月度勾稽\'!{c}$2:{c}$97)'
        emp.append([ident,name,dept,'='+lookup('G'),f'=SUMIF(\'年度底稿\'!A$2:A$9,A{r},\'年度底稿\'!Q$2:Q$9)',f'=D{r}-E{r}','='+lookup('D'),'='+lookup('F'),f'=H{r}-D{r}',f'=D{r}/12','=MIN('+','.join(monthly)+')','=MAX('+','.join(monthly)+')',f'=L{r}-K{r}','=STDEVP('+','.join(monthly)+')',f'=N{r}/J{r}'])
    department=w.create_sheet('部门年度');department.append(['部门','人数','原实发合计','人均原实发年度','员工年度中位数','情景总额','情景增额','重算应发'])
    for r,name in enumerate(['技术部','财务部','销售部','人事部','行政部'],2):
        ids=[row[0] for row in src['基础信息表'].iter_rows(min_row=2,max_col=4,values_only=True) if row[2]==name]
        median_terms=[f'SUMIF(\'员工年度\'!A$2:A$9,{ident},\'员工年度\'!D$2:D$9)' for ident in ids]
        department.append([name,f'=COUNTIF(\'员工年度\'!C$2:C$9,A{r})',f'=SUMIF(\'员工年度\'!C$2:C$9,A{r},\'员工年度\'!D$2:D$9)',f'=C{r}/B{r}','=MEDIAN('+','.join(median_terms)+')',f'=SUMIF(\'员工年度\'!C$2:C$9,A{r},\'员工年度\'!H$2:H$9)',f'=F{r}-C{r}',f'=SUMIF(\'员工年度\'!C$2:C$9,A{r},\'员工年度\'!G$2:G$9)'])
    s=w.create_sheet('审查指标');s.append(['指标','数值'])
    for label,f in [
        ('员工数',"COUNTA('员工年度'!A2:A9)"),('月度记录数',"COUNTA('月度底稿'!B2:B97)"),('原实发年度总额',"SUM('员工年度'!D2:D9)"),('重算应发年度总额',"SUM('员工年度'!G2:G9)"),('个税列年度合计',"SUM('月度底稿'!T2:T97)"),('单次扣款情景年度总额',"SUM('员工年度'!H2:H9)"),('情景差额',"SUM('员工年度'!I2:I9)"),
        ('年度勾稽不一致人数','COUNTIF(\'员工年度\'!F2:F9,">0.005")+COUNTIF(\'员工年度\'!F2:F9,"<-0.005")'),('正情景差额记录数','COUNTIF(\'月度勾稽\'!H2:H97,">0.005")')]:s.append([label,'='+f])
    notes=w.create_sheet('审计结论');notes.append(['主题','源证据','数值影响','核实动作'])
    for row in [
        ['年度勾稽','年度应发工资汇总E:P与各月V，8人12个月','8人原年度Q均等于月表V累计，1231227元','年度表与月表一致不代表工资算法无误'],
        ['命名口径','年度应发工资汇总实际取月表税后实发列','重算应发1556364元，与源年度所谓应发1231227不同','分别报告应发/原实发，不改写源表标题充当业务审批'],
        ['个税重复迹象','每月T扣款含S个税；U=N-T；V=U-S','96条情景差额合计78197元，等于列示个税总和','按所有扣款只扣一次的情景复核；不重估税率，不断言真实税法错误或已批准补发'],
        ['付款状态边界','源表W写已发放','仅源记录文字，缺银行凭证','需要财务核对真实付款与审批，不能视作已补发证据']]:notes.append(row)
    chart=BarChart();chart.title='部门原实发与单次扣款情景（元）'
    for c in [3,6]:chart.add_data(Reference(department,min_col=c,min_row=1,max_row=6),titles_from_data=True)
    chart.set_categories(Reference(department,min_col=1,min_row=2,max_row=6));chart.width=24;chart.height=12;department.add_chart(chart,'A10')
    for s in w:
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
        for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='24476B')
        for row in s.iter_rows(min_row=2):
            for c in row:
                c.alignment=Alignment(wrap_text=True,vertical='top')
                if c.data_type=='f' or isinstance(c.value,(int,float)):c.number_format='#,##0.00'
        for c in range(1,s.max_column+1):s.column_dimensions[col(c)].width=18
        if s.title=='月度底稿':
            for c in range(1,26):s.column_dimensions[col(c)].width=12
            s.column_dimensions['Z'].width=48
            for r in range(2,98):s.row_dimensions[r].height=30
        if s.title=='年度底稿':s.column_dimensions['S'].width=48
        if s.title=='审计结论':
            for c in 'BCD':s.column_dimensions[c].width=52
            for r in range(2,6):s.row_dimensions[r].height=70
        s.sheet_properties.pageSetUpPr.fitToPage=True;s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0;s.print_title_rows='1:1'
    a.output_dir.mkdir(parents=True,exist_ok=True);w.save(a.output_dir/'annual_payroll_audit.xlsx')
if __name__=='__main__':main()
