"""年度工资题独立Judge：原始96记录勾稽、员工/部门统计、变更复算。"""
import argparse,json,math,statistics,subprocess,tempfile,zipfile
from pathlib import Path
from openpyxl import load_workbook
MONTHS=['一月','二月','三月','四月','五月','六月','七月','八月','九月','十月','十一月','十二月']
def near(a,b):return isinstance(a,(int,float)) and not isinstance(a,bool) and math.isclose(a,b,rel_tol=1e-8,abs_tol=.005)
def recalc(f,root,name):
    d=root/name;d.mkdir()
    r=subprocess.run(['libreoffice',f'-env:UserInstallation={(root/(name+"-profile")).as_uri()}','--headless','--convert-to','xlsx','--outdir',str(d),str(f)],capture_output=True,timeout=120)
    if r.returncode or not (d/f.name).exists():raise RuntimeError('LibreOffice重算失败')
    return load_workbook(d/f.name,data_only=True)
def main():
    p=argparse.ArgumentParser();p.add_argument('--artifact-dir',type=Path,required=True);p.add_argument('--result',type=Path,required=True);p.add_argument('--reward-file',type=Path,required=True);a=p.parse_args()
    tests=Path(__file__).parent;inp=tests.parent/'data/input_files';rub=json.loads((tests/'rubric.json').read_text());out=a.artifact_dir/'annual_payroll_audit.xlsx'
    scores={c['id']:0. for c in rub['criteria']};ev={}
    def mark(k,x,e):scores[k]=sum(bool(v) for v in x)/len(x) if x else 0.;ev[k]=e
    w=None
    try:
        if out.exists():w=load_workbook(out)
    except (OSError,ValueError,KeyError,zipfile.BadZipFile):pass
    if w:
        mark('R001',['月度底稿' in w and w['月度底稿'].max_row>1],'原生且有业务底稿')
        source=load_workbook(inp/'完整年度工资管理系统_2019.xlsx',data_only=True)
        records=[]
        for m,name in enumerate(MONTHS,1):
            for r in range(2,10):records.append([m]+[source[name].cell(r,c).value for c in range(1,25)])
        annual=[[source['年度应发工资汇总'].cell(r,c).value for c in range(1,19)] for r in range(2,10)]
        master=[list(r) for r in source['基础信息表'].iter_rows(min_row=2,max_col=4,values_only=True)]
        def raw(sheet,rows,source_col):
            if sheet not in w:return [False]
            s=w[sheet];checks=[s.max_row==len(rows)+1]
            for i,row in enumerate(rows,2):
                for c,value in enumerate(row,1):
                    got=s.cell(i,c).value;checks.append(near(got,value) if isinstance(value,(int,float)) else got==value)
                checks.append('.xlsx!' in str(s.cell(i,source_col).value))
            return checks
        mark('R002',raw('月度底稿',records,26),'96条含原值与状态，不含合计、不复制无关身份资料')
        mark('R003',raw('年度底稿',annual,19),'8人原年度12个月、合计、月均不改写')
        # records 第0列为新增月份，其余下标与原Excel列号一致。
        def derive(rows):
            details=[]
            for r in rows:
                gross=sum(r[7:14]);deduct=sum(r[15:20]);net=gross-deduct
                details.append([r[0],r[1],r[2],gross,deduct,net,r[22],net-r[22],r[14]-gross,r[20]-deduct,r[21]-net,r[22]-(r[21]-r[19])])
            people=[]
            for ident,name,dept,_ in master:
                subset=[r for r in details if r[1]==ident];paid=[r[6] for r in subset];old=next(r[16] for r in annual if r[0]==ident);total=sum(paid);gross=sum(r[3] for r in subset);net=sum(r[5] for r in subset);mean=statistics.mean(paid);std=statistics.pstdev(paid)
                people.append([ident,name,dept,total,old,total-old,gross,net,net-total,mean,min(paid),max(paid),max(paid)-min(paid),std,std/mean])
            departments=[]
            for dept in ['技术部','财务部','销售部','人事部','行政部']:
                ps=[r for r in people if r[2]==dept];paid=[r[3] for r in ps];net=sum(r[7] for r in ps)
                departments.append([dept,len(ps),sum(paid),statistics.mean(paid),statistics.median(paid),net,net-sum(paid),sum(r[6] for r in ps)])
            metrics=[len(master),len(rows),sum(r[6] for r in details),sum(r[3] for r in details),sum(r[19] for r in rows),sum(r[5] for r in details),sum(r[7] for r in details),sum(abs(r[5])>.005 for r in people),sum(r[7]>.005 for r in details)]
            return details,people,departments,metrics
        details,people,depts,metrics=derive(records)
        def numeric(book,sheet,rows):
            if sheet not in book:return [False]
            s=book[sheet];checks=[s.max_row==len(rows)+1]
            for i,row in enumerate(rows,2):
                for c,value in enumerate(row,1):
                    got=s.cell(i,c).value;checks.append(near(got,value) if isinstance(value,(int,float)) else got==value)
            return checks
        required=['月度底稿','年度底稿','月度勾稽','员工年度','部门年度','审查指标','审计结论']
        with tempfile.TemporaryDirectory(prefix='annual-payroll-judge-') as td:
            temp=Path(td)
            try:v=recalc(out,temp,'base')
            except (RuntimeError,subprocess.TimeoutExpired) as e:
                a.result.parent.mkdir(parents=True,exist_ok=True);a.result.write_text(json.dumps({'status':'infrastructure_failure','error':str(e)},ensure_ascii=False));raise SystemExit(2)
            mark('R004',numeric(v,'月度勾稽',details),'每月员工逐笔原值与应发、扣款、净额、算术残差')
            mark('R005',numeric(v,'员工年度',people),'按员工汇总；原实发的极差、总体标准差、变异系数')
            mark('R006',numeric(v,'部门年度',depts),'5部门人数、人均年度和员工年度中位数，非月人混合样本')
            mark('R007',[near(v['审查指标'].cell(i,2).value,x) if '审查指标' in v else False for i,x in enumerate(metrics,2)],'九项跨页总控与异常计数')
            coords=[('月度勾稽',f'{c}{r}') for r in range(2,98) for c in 'DEFGHIJKL']+[('员工年度',f'{c}{r}') for r in range(2,10) for c in 'DEFGHIJKLMNO']+[('部门年度',f'{c}{r}') for r in range(2,7) for c in 'BCDEFGH']+[('审查指标',f'B{r}') for r in range(2,11)]
            mark('R008',[s in w and w[s][c].data_type=='f' for s,c in coords],'所有约定计算保持原生公式，接受等价表达')
            probes=[]
            if all(s in w for s in required) and near(w['月度底稿']['J2'].value,7911):
                q=load_workbook(out);s=q['月度底稿'];s['J2']=12911;s['T2']=1210;s['W2']=15531
                changed=[r[:] for r in records];changed[0][9]+=5000;changed[0][19]+=100;changed[0][22]+=200
                left=[s.cell(2,c).value for c in range(1,27)];right=[s.cell(9,c).value for c in range(1,27)]
                for c in range(1,27):
                    s.cell(2,c).value=right[c-1];s.cell(9,c).value=left[c-1]
                f=temp/'changed.xlsx';q.save(f)
                try:pv=recalc(f,temp,'changed')
                except (RuntimeError,subprocess.TimeoutExpired) as e:
                    a.result.parent.mkdir(parents=True,exist_ok=True);a.result.write_text(json.dumps({'status':'infrastructure_failure','error':str(e)},ensure_ascii=False));raise SystemExit(2)
                nd,np,ng,nm=derive(changed)
                for sheet,old,new in [('月度勾稽',details,nd),('员工年度',people,np),('部门年度',depts,ng)]:
                    for i,(before,after) in enumerate(zip(old,new),2):
                        for c,(b,x) in enumerate(zip(before,after),1):
                            if isinstance(x,(int,float)) and not near(b,x):probes.append(near(pv[sheet].cell(i,c).value,x))
                probes.extend(near(pv['审查指标'].cell(i+2,2).value,x) for i,x in enumerate(nm) if not near(metrics[i],x))
            mark('R009',probes,'绩效/税/原实发变更和首月员工行重排；只奖励应变化结果')
        lines=[' '.join(str(x or '') for x in row) for row in w['审计结论'].iter_rows(min_row=2,max_col=4,values_only=True)] if '审计结论' in w else []
        claims=[['年度','1231227','一致'],['应发','1556364','实发'],['个税','78197','情景'],['已发放','凭证']]
        mark('R010',[any(all(t in line for t in tokens) and len(line)>35 for line in lines) for tokens in claims]+[any(s._charts for s in w),not w._external_links,all(s in w and w[s].freeze_panes and w[s].auto_filter.ref for s in required)],'四类客观证据边界与工作簿可用性，不判断税法或真实付款')
    reward=round(sum(c['weight']*scores[c['id']] for c in rub['criteria']),6)
    if scores['R001']<1:reward=0
    result={'task_id':rub['task_id'],'rubric_version':rub['rubric_version'],'judge_version':'annual-payroll-ledger-v1','reward':reward,'passed':reward>=rub['pass_threshold'],'criterion_results':[{'criterion_id':k,'score':round(v,6),'evidence':ev.get(k,'未满足'),'method':'deterministic'} for k,v in scores.items()]}
    a.result.parent.mkdir(parents=True,exist_ok=True);a.result.write_text(json.dumps(result,ensure_ascii=False,indent=2));a.reward_file.parent.mkdir(parents=True,exist_ok=True);a.reward_file.write_text(str(reward)+'\n')
    raise SystemExit(0 if result['passed'] else 1)
if __name__=='__main__':main()
