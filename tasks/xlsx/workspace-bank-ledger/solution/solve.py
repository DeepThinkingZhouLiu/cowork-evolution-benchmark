#!/usr/bin/env python3
import argparse, re
from pathlib import Path
from collections import Counter, defaultdict
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

def num(v):
    if isinstance(v,(int,float)): return float(v)
    s=str(v or '').replace(',','').replace('元','').strip()
    return float(s) if s else 0.0
def money(v): return round(float(v)+0.00000001,2)
def style(ws, widths=None):
    ws.freeze_panes='A2'; ws.sheet_view.showGridLines=False
    if ws.max_row>=1:
        ws.auto_filter.ref=f'A1:{chr(64+ws.max_column)}{ws.max_row}'
        for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F4E78'); c.alignment=Alignment(horizontal='center',wrap_text=True)
    if widths:
        for col,w in widths.items(): ws.column_dimensions[col].width=w
def table(ws,name):
    if ws.max_row>1:
        t=Table(displayName=name,ref=f'A1:{chr(64+ws.max_column)}{ws.max_row}'); t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True); ws.add_table(t)
def extract(src):
    w=load_workbook(src,data_only=True); home={str(r[0]).strip():r[1] for r in w['主页'].iter_rows(min_row=2,values_only=True) if r[0]}
    months=[]; cp=Counter(); sums=defaultdict(lambda:[0.0,0.0])
    prev=num(home.get('初始余额'))
    for ws in w.worksheets[1:]:
        m=ws.title.strip(); debit=credit=0.0; ending=None
        for row in ws.iter_rows(min_row=5,values_only=True):
            if row[2]=='本月合计': ending=num(row[8]); break
            if not row[2] or not row[5]: continue
            d=num(row[6]); c=num(row[7]); debit+=d; credit+=c; cp[str(row[5]).strip()]+=1; sums[str(row[5]).strip()][0]+=d; sums[str(row[5]).strip()][1]+=c
        if ending is None: ending=money(prev+debit-credit)
        months.append((m,money(prev),money(debit),money(credit),money(ending))); prev=ending
    return home,months,cp,sums
def build(input_dir,output_dir):
    output_dir.mkdir(parents=True,exist_ok=True); src=input_dir/'银行存款日记账_完整版.xlsx'; home,months,counts,sums=extract(src)
    wb=load_workbook(src); wb.remove(wb['主页'])
    # Remove source monthly tabs and rebuild a compact analytical workbook.
    for ws in list(wb.worksheets): wb.remove(ws)
    ws=wb.create_sheet('月度资金趋势'); ws.append(['月份','期初余额','本月借方','本月贷方','月末结存','本年累计借方','本年累计贷方'])
    for i,(m,opening,d,c,end) in enumerate(months,2):
        ws.append([m,opening,d,c,end, f'=SUM($C$2:C{i})', f'=SUM($D$2:D{i})'])
        for col in range(2,8): ws.cell(i,col).number_format='#,##0.00;[Red]-#,##0.00'
    style(ws,{'A':12,'B':16,'C':16,'D':16,'E':16,'F':18,'G':18}); table(ws,'MonthlyTrendTable')
    ws=wb.create_sheet('年度收支汇总'); ws.append(['指标','数值']); total_d=money(sum(x[2] for x in months)); total_c=money(sum(x[3] for x in months)); ending=months[-1][4]; target=num(home.get('年度目标余额')); opening=num(home.get('初始余额'))
    for k,v in [('年度累计借方',total_d),('年度累计贷方',total_c),('年度收支净额',money(total_d-total_c)),('年末结存余额',ending),('年度目标余额',target),('年度目标完成率',ending/target if target else 0)]: ws.append([k,v])
    for r in range(2,ws.max_row+1): ws.cell(r,2).number_format='0.00%' if r==7 else '#,##0.00;[Red]-#,##0.00'
    style(ws,{'A':24,'B':22}); table(ws,'AnnualSummaryTable')
    ws=wb.create_sheet('账户基础信息'); ws.append(['项目','详情']);
    for k,v in home.items(): ws.append([k,v])
    style(ws,{'A':22,'B':32}); table(ws,'AccountInfoTable')
    ws=wb.create_sheet('高频交易对手'); ws.append(['排名','交易对手','交易次数','借方合计','贷方合计','交易总额'])
    for rank,(entity,n) in enumerate(counts.most_common(),1): ws.append([rank,entity,n,money(sums[entity][0]),money(sums[entity][1]),f'=D{rank+1}+E{rank+1}'])
    for r in range(2,ws.max_row+1):
        for c in [4,5,6]: ws.cell(r,c).number_format='#,##0.00;[Red]-#,##0.00'
    style(ws,{'A':10,'B':28,'C':12,'D':16,'E':16,'F':16}); table(ws,'CounterpartyTable')
    ws=wb.create_sheet('分析报告'); ws.append(['章节','分析内容'])
    maxd=max(months,key=lambda x:x[2]); maxc=max(months,key=lambda x:x[3]); minb=min(months,key=lambda x:x[4]); maxb=max(months,key=lambda x:x[4]); h1=sum(x[2]-x[3] for x in months[:6]); h2=sum(x[2]-x[3] for x in months[6:])
    text=[('账户基础信息',f"开户银行：{home.get('开户银行','')}；银行账号：{home.get('银行账号','')}；启用日期：{home.get('启用日期','')}；初始余额：{home.get('初始余额','')}；年度目标余额：{home.get('年度目标余额','')}。"),('年度收支汇总',f'全年借方 {total_d:,.2f} 元，贷方 {total_c:,.2f} 元，净额 {total_d-total_c:,.2f} 元，年末结存 {ending:,.2f} 元，目标完成率 {ending/target:.2%}。'),('月度余额变化分析',f'借方最高为 {maxd[0]}（{maxd[2]:,.2f} 元），贷方最高为 {maxc[0]}（{maxc[3]:,.2f} 元）；最低余额为 {minb[0]}（{minb[4]:,.2f} 元），最高余额为 {maxb[0]}（{maxb[4]:,.2f} 元）。上半年净变动 {h1:,.2f} 元，下半年净变动 {h2:,.2f} 元。'),('高频交易对手方统计','、'.join(f'{e}（{n}次）' for e,n in counts.most_common(10))),('结论与优化建议','1. 对负余额月份设置阶段性流动性预警和最低现金缓冲；2. 对高频交易对手按月核对并维护付款计划；3. 将可用余额与短期支出预测联动，减少资金闲置；4. 每月复核借贷方向、余额和异常交易。')]
    for row in text: ws.append(row)
    style(ws,{'A':24,'B':110}); table(ws,'AnalysisReportTable')
    ws=wb.create_sheet('Source Notes'); ws.append(['项目','内容']); ws.append(['来源','Workspace-Bench-Lite task 255；Workspace-Bench-Workspaces revision e245d63bfa20cfdb708cd8e78145ffb087155857']); ws.append(['口径','月度及年度金额由可见交易明细重新求和；metadata rubric 与可见输入数值不一致时以输入为准']); ws.append(['状态','pending_human_review / uncalibrated'])
    style(ws,{'A':20,'B':110}); table(ws,'SourceNotesTable')
    wb.calculation.fullCalcOnLoad=True; wb.calculation.forceFullCalc=True
    out=output_dir/'月度资金流动趋势表.xlsx'; wb.save(out); return out
if __name__=='__main__':
    ap=argparse.ArgumentParser(); ap.add_argument('--input-dir',type=Path,required=True); ap.add_argument('--output-dir',type=Path,required=True); a=ap.parse_args(); build(a.input_dir,a.output_dir)
