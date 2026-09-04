#!/usr/bin/env python3
import argparse, hashlib, json, re, zipfile
from pathlib import Path
from openpyxl import load_workbook

def num(v):
    if isinstance(v,(int,float)): return float(v)
    return float(str(v or '0').replace(',','').replace('元',''))
def digest(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def c(i,ok,e): return {'criterion_id':f'R{i:03d}','score':1.0 if ok else 0.0,'evidence':e,'method':'deterministic'}
def truth(src):
    w=load_workbook(src,data_only=True); home={str(r[0]).strip():r[1] for r in w['主页'].iter_rows(min_row=2,values_only=True) if r[0]}; ms=[]; counts={}; sums={}; prev=num(home['初始余额'])
    for s in w.worksheets[1:]:
        d=cr=0.; end=None
        for row in s.iter_rows(min_row=5,values_only=True):
            if row[2]=='本月合计': end=num(row[8]); break
            if not row[2] or not row[5]: continue
            d+=num(row[6]); cr+=num(row[7]); ent=str(row[5]).strip(); counts[ent]=counts.get(ent,0)+1; sums.setdefault(ent,[0.,0.]); sums[ent][0]+=num(row[6]); sums[ent][1]+=num(row[7])
        ms.append((s.title.strip(),round(prev,2),round(d,2),round(cr,2),round(end if end is not None else prev+d-cr,2))); prev=end if end is not None else prev+d-cr
    return home,ms,counts,sums
def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--artifact-dir',type=Path,required=True); ap.add_argument('--result',type=Path,required=True); ap.add_argument('--reward-file',type=Path,required=True); a=ap.parse_args(); root=Path(__file__).parents[1]; src=root/'data/input_files/银行存款日记账_完整版.xlsx'; out=a.artifact_dir/'月度资金流动趋势表.xlsx'; rs=[]
    native=out.is_file() and zipfile.is_zipfile(out); rs.append(c(1,native,'输出存在且为有效 XLSX'))
    if native:
      try:
        home,ms,counts,sums=truth(src); w=load_workbook(out,data_only=True); trend=w['月度资金趋势']; rows=list(trend.iter_rows(min_row=2,values_only=True)); totald=round(sum(x[2] for x in ms),2); totalc=round(sum(x[3] for x in ms),2); end=ms[-1][4]; target=num(home['年度目标余额']); annual=w['年度收支汇总']; amap={r[0]:r[1] for r in annual.iter_rows(min_row=2,values_only=True)}; report=' '.join(str(r[1] or '') for r in w['分析报告'].iter_rows(min_row=2,values_only=True)); top=sorted(counts.items(),key=lambda x:(-x[1],x[0]))
        top5=dict(top[:5]); report_sections=[r[0] for r in w['分析报告'].iter_rows(min_row=2,values_only=True)]
        rs += [c(2,len(rows)==12,'趋势表12个月'), c(3,[r[0] for r in rows]==[x[0] for x in ms],'月份顺序'), c(4,all(abs(float(rows[i][2])-ms[i][2])<.01 and abs(float(rows[i][3])-ms[i][3])<.01 and abs(float(rows[i][4])-ms[i][4])<.01 for i in range(12)),'月度借贷与余额从明细重算一致'), c(5,abs(num(amap.get('年度累计借方'))-totald)<.01 and abs(num(amap.get('年度累计贷方'))-totalc)<.01,'年度累计从输入重算'), c(6,abs(num(amap.get('年度收支净额'))-(totald-totalc))<.01 and abs(num(amap.get('年末结存余额'))-end)<.01,'净额与年末余额一致'), c(7,abs(num(amap.get('年度目标完成率'))-end/target)<.0001,'目标完成率一致'), c(8,top5.get('税务局')==19 and top5.get('应付职工薪酬')==18 and top5.get('武汉XX建材有限公司')==17 and top5.get('杭州XX服装有限公司')==13 and top5.get('上海XX贸易有限公司')==13,'Top交易对手频次正确'), c(9,'账户基础信息' in w.sheetnames and str(home['开户银行']) in ' '.join(str(r[1]) for r in w['账户基础信息'].iter_rows(min_row=2,values_only=True)),'账户基础信息保留'), c(10,all(x in report_sections for x in ['账户基础信息','年度收支汇总','月度余额变化分析','高频交易对手方统计','结论与优化建议']),'分析报告章节完整'), c(11,'流动性' in report and '交易对手' in report and '资金' in report,'建议覆盖业务风险'), c(12,'月度资金趋势' in w.sheetnames and '高频交易对手' in w.sheetnames,'分析工作表存在'), c(13,all((a.artifact_dir/x).is_file() for x in ['月度资金流动趋势表.xlsx']) and len(list(a.artifact_dir.iterdir()))==1,'无额外输出')]
      except Exception as e: rs += [c(i,False,f'解析失败: {e}') for i in range(2,14)]
    else: rs += [c(i,False,'输出缺失或损坏') for i in range(2,14)]
    reward=sum(x['score'] for x in rs)/len(rs); a.result.parent.mkdir(parents=True,exist_ok=True); a.result.write_text(json.dumps({'task_id':'cowork-evo/workspace-bank-ledger','rubric_version':'workspace-bank-ledger-v0.1','criteria':rs,'reward':reward},ensure_ascii=False,indent=2)); a.reward_file.parent.mkdir(parents=True,exist_ok=True); a.reward_file.write_text(f'{reward:.6f}\n')
if __name__=='__main__': main()
