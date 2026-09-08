"""只检验食堂这一题：不能用批量题目生成器代替逐题阅读。"""
import argparse,json,shutil,subprocess,sys
from pathlib import Path
from openpyxl import load_workbook,Workbook
def main():
    p=argparse.ArgumentParser();p.add_argument('--output',type=Path,required=True);a=p.parse_args();out=a.output.resolve();out.mkdir(parents=True,exist_ok=False);root=Path(__file__).resolve().parents[1]
    fresh=out/'fresh';subprocess.run([sys.executable,str(root/'solution/solve.py'),'--input-dir',str(root/'data/input_files'),'--output-dir',str(fresh)],check=True)
    cases={'fresh':fresh,'golden':root/'solution/Golden'}
    for label in ['empty','blank','corrupt','unchanged','no_rounding','invented_compliance','mixed_period','wrong_inventory','constant_formulas']:
        folder=out/label;folder.mkdir();cases[label]=folder;target=folder/'canteen_review.xlsx'
        if label=='empty':continue
        if label=='blank':Workbook().save(target);continue
        if label=='corrupt':target.write_bytes(b'broken');continue
        if label=='unchanged':shutil.copy2(root/'data/input_files/2-食堂菜品出入库表.xlsx',target);continue
        w=load_workbook(fresh/'canteen_review.xlsx')
        if label=='no_rounding':
            for r in range(2,10):w['采购计划'][f'F{r}']=f'=MAX(0,D{r}+E{r}-C{r})'
        if label=='invented_compliance':
            for r in range(2,8):
                for c in range(2,5):w['证据审查'].cell(r,c,'全部合格，完成率100%')
            for r in range(2,9):w['7S检查计划'].cell(r,5,'已完成')
        if label=='mixed_period':w['成本情景']['G2']="=SUM(B2:B5)+SUM('库存底稿'!J2:J9)"
        if label=='wrong_inventory':w['库存底稿']['B2']=80
        if label=='constant_formulas':
            cached=out/'cached';cached.mkdir()
            subprocess.run(['libreoffice',f'-env:UserInstallation={(out/"cache-profile").as_uri()}','--headless','--convert-to','xlsx','--outdir',str(cached),str(fresh/'canteen_review.xlsx')],check=True,capture_output=True)
            values=load_workbook(cached/'canteen_review.xlsx',data_only=True)
            for s in w:
                for row in s:
                    for c in row:
                        if c.data_type=='f':c.value='='+str(values[s.title][c.coordinate].value)
        w.save(target)
    scores={}
    for label,folder in cases.items():
        j=out/'judges'/label;j.mkdir(parents=True)
        with (j/'judge.log').open('w') as log:r=subprocess.run([sys.executable,str(root/'tests/judge.py'),'--artifact-dir',str(folder),'--result',str(j/'result.json'),'--reward-file',str(j/'reward.txt')],stdout=log,stderr=subprocess.STDOUT,timeout=300)
        assert r.returncode in (0,1),(label,r.returncode)
        scores[label]=json.loads((j/'result.json').read_text())['reward'];print(label,scores[label],flush=True)
    assert scores['fresh']==scores['golden']==1
    assert all(scores[x]==0 for x in ['empty','blank','corrupt','unchanged'])
    assert all(0<scores[x]<1 for x in ['no_rounding','invented_compliance','mixed_period','wrong_inventory'])
    assert scores['constant_formulas']<.8
    repeats=[]
    for n in range(3):
        j=out/f'repeat-{n}';j.mkdir()
        subprocess.run([sys.executable,str(root/'tests/judge.py'),'--artifact-dir',str(fresh),'--result',str(j/'result.json'),'--reward-file',str(j/'reward.txt')],capture_output=True,check=True,timeout=300);repeats.append(json.loads((j/'result.json').read_text())['reward'])
    assert repeats==[1]*3
    (out/'validation.json').write_text(json.dumps({'task':'workspace-canteen-stock-procurement','scores':scores,'repeated_rewards':repeats,'all_assertions_passed':True},ensure_ascii=False,indent=2))
if __name__=='__main__':main()

