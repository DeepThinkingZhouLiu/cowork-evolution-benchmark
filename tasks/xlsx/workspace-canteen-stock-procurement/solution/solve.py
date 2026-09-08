"""食堂题专属参考解：源证据与假设分离，按菜品关联采购。"""
import argparse
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.utils.datetime import from_excel
from openpyxl.chart import BarChart,Reference
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
def main():
    p=argparse.ArgumentParser();p.add_argument('--input-dir',type=Path,required=True);p.add_argument('--output-dir',type=Path,required=True);a=p.parse_args()
    w=Workbook();w.remove(w.active)
    s=w.create_sheet('费用底稿');s.append(['日期','类型','细项','原单位','数量','单价元','原金额元','重算金额元','来源'])
    src=load_workbook(a.input_dir/'3-食堂费用支出统计表.xlsx',data_only=True).active
    for r in range(7,15):
        n=s.max_row+1;s.append([from_excel(src.cell(r,3).value)]+[src.cell(r,c).value for c in range(4,10)]+[f'=E{n}*F{n}',f'3-食堂费用支出统计表.xlsx!Sheet1!C{r}:I{r}']);s.cell(n,1).number_format='yyyy-mm-dd'
    s=w.create_sheet('库存底稿');s.append(['菜品','到货kg','到货单价','出货kg','出货单价','原到货元','原出货元','原结存kg','原结存元','重算到货元','重算出货元','重算结存kg','重算结存元','来源'])
    src=load_workbook(a.input_dir/'2-食堂菜品出入库表.xlsx',data_only=True).active
    for r in range(10,18):
        n=s.max_row+1;s.append([src.cell(r,c).value for c in (6,7,8,13,14,9,15,16,17)]+[f'=B{n}*C{n}',f'=D{n}*E{n}',f'=B{n}-D{n}',f'=J{n}-K{n}',f'2-食堂菜品出入库表.xlsx!{src.title}!F{r}:Q{r}'])
    s=w.create_sheet('参数');s.append(['情景假设','值']);s.append(['需求倍数',1.2]);s.append(['安全库存率',.1]);s.append(['调味品单价涨幅',.15])
    s=w.create_sheet('采购计划');s.append(['菜品','基期出货kg','现存kg','情景消耗kg','安全库存kg','需采购整kg','采购单价','采购金额元','预计期末kg'])
    for r in range(2,10):
        name=w['库存底稿'].cell(r,1).value
        s.append([name,f'=SUMIF(\'库存底稿\'!A$2:A$9,A{r},\'库存底稿\'!D$2:D$9)',f'=SUMIF(\'库存底稿\'!A$2:A$9,A{r},\'库存底稿\'!L$2:L$9)',f'=B{r}*\'参数\'!B$2',f'=D{r}*\'参数\'!B$3',f'=MAX(0,ROUNDUP(D{r}+E{r}-C{r},0))',f'=SUMIF(\'库存底稿\'!A$2:A$9,A{r},\'库存底稿\'!C$2:C$9)',f'=F{r}*G{r}',f'=C{r}+F{r}-D{r}'])
    s=w.create_sheet('成本情景');s.append(['费用类型','基期元','情景元','增加元',None,'指标','结果'])
    for r,kind in enumerate(['食品支出','调味品','房租设备','人工成本'],2):
        s.cell(r,1,kind);s.cell(r,2,f'=SUMIF(\'费用底稿\'!B$2:B$9,A{r},\'费用底稿\'!H$2:H$9)');s.cell(r,3,f'=IF(A{r}="调味品",B{r}*(1+\'参数\'!B$4),B{r})');s.cell(r,4,f'=C{r}-B{r}')
    metrics=[('费用合计','SUM(B2:B5)'),('调价后合计','SUM(C2:C5)'),('增加额','SUM(D2:D5)'),('库存到货金额',"SUM('库存底稿'!J2:J9)"),('出货金额',"SUM('库存底稿'!K2:K9)"),('结存金额',"SUM('库存底稿'!M2:M9)"),('结存金额占到货比例','G7/G5'),('到货kg',"SUM('库存底稿'!B2:B9)"),('出货kg',"SUM('库存底稿'!D2:D9)"),('结存kg',"SUM('库存底稿'!L2:L9)"),('采购kg',"SUM('采购计划'!F2:F9)"),('采购金额',"SUM('采购计划'!H2:H9)"),('预计期末kg',"SUM('采购计划'!I2:I9)")]
    for r,(name,f) in enumerate(metrics,2):s.cell(r,6,name);s.cell(r,7,'='+f)
    s=w.create_sheet('证据审查');s.append(['主题','判断','现有证据','尚缺材料/下一步'])
    records=[
        ['期间','不可合并同月成本','费用序号44562起为2022-01；库存20xx年11月','需要完整年份和相同期间；当前分别核算'],
        ['单位','保留异常待核实','电费单位写公斤；兼职费用出现个/人','不得直接当作电量或人数；取账单与计量定义核验'],
        ['人均餐费','无法确定','套表的中晚夜宵人数为空，费用表未给就餐人数','补真实用餐人数与同期费用，不能编造实际人均'],
        ['健康证明','证据不足','指南要求持健康合格证和每年体检，未附证书/体检记录','索取工作人员有效证书及体检日期，不填全部合格'],
        ['培训','证据不足','7S要求入职不少于2小时、月度参与率不低于90%；无签到成绩','索取培训计划、签到和成绩，不填完成率'],
        ['保质期','证据不足','库存只有数量价格与领取人，无批次/有效期','索取批次保质期和处理记录，不能由结存正数断言未过期']]
    for row in records:s.append(row)
    s=w.create_sheet('7S检查计划');s.append(['维度','标准控制','频率/阈值','需查证据','状态'])
    plan=[
        ['整理','清查过期/变质食材与调料','每周一次','保质期记录、处置日期数量方式责任人'],
        ['整顿','按类别和生熟分开存放并标识','现场检查','货架标签、生熟区域和存放实况'],
        ['清扫','厨房餐后清洁、下班深度消毒','每餐及每日','清洁消毒记录与现场检查'],
        ['清洁','内部审核并跟踪整改','每周一次','审核报告、不符合项和整改闭环证据'],
        ['素养','入职7S培训和考核','不少于2小时','培训计划、签到、教材、试卷成绩'],
        ['安全','燃气泄漏安全检查','每月不少于一次','安检记录、检测数据、检查人日期'],
        ['节约','依据消耗分析调整采购','每周分析','采购计划、实际采购、库存和消耗分析']]
    for row in plan:s.append(row+['待核验'])
    chart=BarChart();chart.title='情景采购金额（元）';chart.add_data(Reference(w['采购计划'],min_col=8,min_row=1,max_row=9),titles_from_data=True);chart.set_categories(Reference(w['采购计划'],min_col=1,min_row=2,max_row=9));chart.width=22;chart.height=10;w['采购计划'].add_chart(chart,'A12')
    for s in w:
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
        for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='24476B')
        for row in s.iter_rows(min_row=2):
            for c in row:
                c.alignment=Alignment(wrap_text=True,vertical='top')
                if c.data_type=='f' or isinstance(c.value,(int,float)):c.number_format='#,##0.00'
        for c in range(1,s.max_column+1):s.column_dimensions[get_column_letter(c)].width=21
        if s.title in ['证据审查','7S检查计划']:
            for c in 'BCD':s.column_dimensions[c].width=45
            for r in range(2,s.max_row+1):s.row_dimensions[r].height=65
        s.sheet_properties.pageSetUpPr.fitToPage=True;s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=1;s.print_title_rows='1:1'
    w['成本情景']['G8'].number_format='0.00%'
    a.output_dir.mkdir(parents=True,exist_ok=True);w.save(a.output_dir/'canteen_review.xlsx')
if __name__=='__main__':main()
