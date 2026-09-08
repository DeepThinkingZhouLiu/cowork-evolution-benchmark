"""食堂题独立 Judge；从原费用和库存重算，不导入参考解、不读取 Golden。"""
import argparse,json,math,subprocess,tempfile,zipfile
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from openpyxl.formula.translate import Translator

def near(a,b):return isinstance(a,(int,float)) and not isinstance(a,bool) and math.isclose(a,b,rel_tol=1e-8,abs_tol=.005)
def recalc(path,root,name):
    d=root/name;d.mkdir()
    r=subprocess.run(['libreoffice',f'-env:UserInstallation={(root/(name+"-profile")).as_uri()}','--headless','--convert-to','xlsx','--outdir',str(d),str(path)],capture_output=True,timeout=100)
    if r.returncode or not (d/path.name).exists():raise RuntimeError('LibreOffice重算无产物')
    return load_workbook(d/path.name,data_only=True)
def main():
    p=argparse.ArgumentParser();p.add_argument('--artifact-dir',type=Path,required=True);p.add_argument('--result',type=Path,required=True);p.add_argument('--reward-file',type=Path,required=True);a=p.parse_args()
    tests=Path(__file__).parent;inp=tests.parent/'data/input_files';rub=json.loads((tests/'rubric.json').read_text());out=a.artifact_dir/'canteen_review.xlsx'
    scores={c['id']:0. for c in rub['criteria']};ev={}
    def mark(k,checks,why):scores[k]=sum(bool(v) for v in checks)/len(checks) if checks else 0.;ev[k]=why
    w=None
    try:
        if out.is_file():w=load_workbook(out)
    except (OSError,ValueError,KeyError,zipfile.BadZipFile):pass
    if w:
        mark('R001',['库存底稿' in w and w['库存底稿'].max_row>1],'有效且非空库存业务底稿')
        expense=load_workbook(inp/'3-食堂费用支出统计表.xlsx',data_only=True).active
        erows=[[from_excel(expense.cell(r,3).value)]+[expense.cell(r,c).value for c in range(4,10)] for r in range(7,15)]
        stock=load_workbook(inp/'2-食堂菜品出入库表.xlsx',data_only=True).active
        srows=[[stock.cell(r,c).value for c in (6,7,8,13,14,9,15,16,17)] for r in range(10,18)]
        def raw(sheet,rows,col):
            if sheet not in w:return [False]
            checks=[w[sheet].max_row==len(rows)+1]
            for r,row in enumerate(rows,2):
                for c,value in enumerate(row,1):
                    got=w[sheet].cell(r,c).value;checks.append(near(got,value) if isinstance(value,(int,float)) else got==value)
                checks.append('.xlsx!' in str(w[sheet].cell(r,col).value))
            return checks
        mark('R002',raw('费用底稿',erows,9),'8条实际费用、Excel日期、原单位/原值')
        mark('R003',raw('库存底稿',srows,14),'8菜品9个原字段、来源；不能混入模板行或改原库存')
        parameters=[w['参数'].cell(r,2).value if '参数' in w else None for r in (2,3,4)]
        mark('R004',[near(x,y) for x,y in zip(parameters,[1.2,.1,.15])],'情景参数与题面一致，非声称源事实')
        cats=['食品支出','调味品','房租设备','人工成本']
        expenses={k:sum(r[4]*r[5] for r in erows if r[1]==k) for k in cats}
        def procurement(rows,factor):
            result={}
            for name,qty,price,used,outprice,*_ in rows:
                remaining=qty-used;consume=used*factor;safety=consume*.1;buy=max(0,math.ceil(consume+safety-remaining-1e-9))
                result[name]=[used,remaining,consume,safety,buy,price,buy*price,remaining+buy-consume]
            return result
        plans=procurement(srows,1.2)
        arrivals=sum(r[1]*r[2] for r in srows);dispatch=sum(r[3]*r[4] for r in srows)
        metric=[sum(expenses.values()),sum(expenses.values())+expenses['调味品']*.15,expenses['调味品']*.15,arrivals,dispatch,arrivals-dispatch,(arrivals-dispatch)/arrivals,sum(r[1] for r in srows),sum(r[3] for r in srows),sum(r[1]-r[3] for r in srows),sum(r[4] for r in plans.values()),sum(r[6] for r in plans.values()),sum(r[7] for r in plans.values())]
        required=['费用底稿','库存底稿','参数','采购计划','成本情景','证据审查','7S检查计划']
        with tempfile.TemporaryDirectory(prefix='canteen-judge-') as td:
            temp=Path(td)
            try:v=recalc(out,temp,'base')
            except (RuntimeError,subprocess.TimeoutExpired) as err:
                a.result.parent.mkdir(parents=True,exist_ok=True);a.result.write_text(json.dumps({'status':'infrastructure_failure','error':str(err)},ensure_ascii=False));raise SystemExit(2)
            def plan_checks(book,expected):
                if '采购计划' not in book:return [False]
                seen=[];checks=[]
                for row in book['采购计划'].iter_rows(min_row=2,max_row=9,max_col=9,values_only=True):
                    seen.append(row[0]);checks.extend(near(x,y) for x,y in zip(row[1:],expected.get(row[0],[math.nan]*8)))
                checks.append(set(seen)==set(expected))
                return checks
            mark('R005',plan_checks(v,plans),'按品计算安全库存、逐项向上取整、采购费和期末，不按合计取整')
            checks=[]
            if '成本情景' in v:
                s=v['成本情景']
                for r,k in enumerate(cats,2):
                    e=expenses[k];new=e*(1.15 if k=='调味品' else 1)
                    checks.extend([s.cell(r,1).value==k,near(s.cell(r,2).value,e),near(s.cell(r,3).value,new),near(s.cell(r,4).value,new-e)])
                checks += [near(s.cell(r,7).value,value) for r,value in enumerate(metric,2)]
            if '库存底稿' in v:
                for i,row in enumerate(srows,2):checks.extend(near(v['库存底稿'].cell(i,c).value,value) for c,value in zip(range(10,14),[row[1]*row[2],row[3]*row[4],row[1]-row[3],row[1]*row[2]-row[3]*row[4]]))
            if '费用底稿' in v:checks.extend(near(v['费用底稿'].cell(i,8).value,row[4]*row[5]) for i,row in enumerate(erows,2))
            mark('R006',checks,'基期/情景成本、库存恒等式、单位与汇总，不把不同期间相加')
            coords=[('费用底稿',f'H{r}') for r in range(2,10)]+[('库存底稿',f'{c}{r}') for r in range(2,10) for c in 'JKLM']+[('采购计划',f'{c}{r}') for r in range(2,10) for c in 'BCDEFGHI']+[('成本情景',f'{c}{r}') for r in range(2,6) for c in 'BCD']+[('成本情景',f'G{r}') for r in range(2,15)]
            mark('R007',[s in w and w[s][c].data_type=='f' for s,c in coords],'声明计算均用原生公式')
            checks=[]
            if all(s in w for s in required):
                q=load_workbook(out);q['参数']['B2']=1.4;q['参数']['B4']=.2
                if near(q['库存底稿']['B2'].value,30):
                    q['库存底稿']['B2']=35;modified=[row[:] for row in srows];modified[0][1]=35
                    s=q['库存底稿']
                    left=[s.cell(2,c).value for c in range(1,15)];right=[s.cell(9,c).value for c in range(1,15)]
                    for c in range(1,15):
                        lv,rv=left[c-1],right[c-1]
                        if isinstance(lv,str) and lv.startswith('='):lv=Translator(lv,origin=s.cell(2,c).coordinate).translate_formula(s.cell(9,c).coordinate)
                        if isinstance(rv,str) and rv.startswith('='):rv=Translator(rv,origin=s.cell(9,c).coordinate).translate_formula(s.cell(2,c).coordinate)
                        s.cell(2,c,rv);s.cell(9,c,lv)
                    probe=temp/'changed.xlsx';q.save(probe)
                    try:pv=recalc(probe,temp,'probe')
                    except (RuntimeError,subprocess.TimeoutExpired) as err:
                        a.result.parent.mkdir(parents=True,exist_ok=True);a.result.write_text(json.dumps({'status':'infrastructure_failure','error':str(err)},ensure_ascii=False));raise SystemExit(2)
                    newplans=procurement(modified,1.4)
                    # 只对真正应改变的数值给行为分，固定不变的原单价不产生伪依赖奖励。
                    actualplans={r[0]:r[1:] for r in pv['采购计划'].iter_rows(min_row=2,max_row=9,max_col=9,values_only=True)}
                    for name,values in newplans.items():
                        for k in (2,3,4,6,7):
                            if not near(values[k],plans[name][k]):
                                checks.append(near(actualplans.get(name,[None]*8)[k],values[k]))
                    checks.extend([near(pv['成本情景']['G3'].value,sum(expenses.values())+expenses['调味品']*.2),near(pv['成本情景']['G5'].value,arrivals+9),near(pv['成本情景']['G12'].value,sum(r[4] for r in newplans.values()))])
            mark('R008',checks,'需求倍数、调价幅度、现存与菜品重排后的变更行为')
        lines=[' '.join(str(c or '') for c in row) for row in w['证据审查'].iter_rows(min_row=2,max_col=4,values_only=True)] if '证据审查' in w else []
        issues=[('期间',['2022','20xx']),('单位',['电费','公斤']),('人均',['人数','不能']),('健康',['证书','记录']),('培训',['2小时','90%']),('保质期',['批次','不能'])]
        mark('R009',[any(topic in row and all(t in row for t in words) and len(row)>35 for row in lines) for topic,words in issues],'六项证据边界；仅核对可验证要点，不声称语义全面审查')
        plan_checks_7s=[]
        expected=[('整理',['过期','每周']),('整顿',['生熟','标识']),('清扫',['清洁','消毒']),('清洁',['每周','审核']),('素养',['2小时','成绩']),('安全',['燃气','每月']),('节约',['采购','消耗'])]
        lines=list(w['7S检查计划'].iter_rows(min_row=2,max_row=8,max_col=5,values_only=True)) if '7S检查计划' in w else []
        for name,words in expected:
            plan_checks_7s.append(any(row[0]==name and row[4]=='待核验' and all(word in ' '.join(str(c or '') for c in row) for word in words) for row in lines))
        mark('R010',plan_checks_7s+[any(s._charts for s in w),not w._external_links,all(s in w and w[s].auto_filter.ref and w[s].freeze_panes for s in required)],'逐维度标准/频率/证据/待核验，及原生图表与可用性')
    score=round(sum(c['weight']*scores[c['id']] for c in rub['criteria']),6)
    if scores['R001']<1:score=0
    result={'task_id':rub['task_id'],'rubric_version':rub['rubric_version'],'judge_version':'canteen-stock-procurement-v1','reward':score,'passed':score>=rub['pass_threshold'],'criterion_results':[{'criterion_id':k,'score':round(v,6),'evidence':ev.get(k,'未满足'),'method':'deterministic'} for k,v in scores.items()]}
    a.result.parent.mkdir(parents=True,exist_ok=True);a.result.write_text(json.dumps(result,ensure_ascii=False,indent=2));a.reward_file.parent.mkdir(parents=True,exist_ok=True);a.reward_file.write_text(str(score)+'\n')
    raise SystemExit(0 if result['passed'] else 1)
if __name__=='__main__':main()
