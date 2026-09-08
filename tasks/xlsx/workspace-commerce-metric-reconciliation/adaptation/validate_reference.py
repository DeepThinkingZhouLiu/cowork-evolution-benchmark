"""本题专属负例：跨期、错分母、平均ROI、虚构未知、伪公式。"""
import argparse,json,shutil,subprocess,sys
from pathlib import Path
from openpyxl import Workbook,load_workbook

def main():
    p=argparse.ArgumentParser();p.add_argument('--output',type=Path,required=True);a=p.parse_args()
    out=a.output.resolve();out.mkdir(parents=True,exist_ok=False);root=Path(__file__).resolve().parents[1]
    fresh=out/'fresh';subprocess.run([sys.executable,str(root/'solution/solve.py'),'--input-dir',str(root/'data/input_files'),'--output-dir',str(fresh)],check=True)
    cases={'fresh':fresh,'golden':root/'solution/Golden'}
    for name in ['empty','blank','corrupt','unchanged','full_month','wrong_denominator','average_roi','invent_unknown','overwrite_source','constant_formulas']:
        d=out/name;d.mkdir();cases[name]=d;f=d/'commerce_review.xlsx'
        if name=='empty':continue
        if name=='blank':Workbook().save(f);continue
        if name=='corrupt':f.write_bytes(b'broken');continue
        if name=='unchanged':shutil.copy2(next((root/'data/input_files').glob('*.xlsx')),f);continue
        w=load_workbook(fresh/'commerce_review.xlsx')
        if name=='full_month':w['经营指标']['B2']="=SUM('周报底稿'!B2:B17)"
        if name=='wrong_denominator':
            for r in range(2,9):w['单品净化'][f'E{r}']=f'=G{r}'
        if name=='average_roi':w['经营指标']['B11']="=AVERAGE('周报底稿'!K2:K8)"
        if name=='invent_unknown':
            for r in (19,20,21):w['经营指标'][f'B{r}']=0
        if name=='overwrite_source':w['单品底稿']['J2']=49/1314
        if name=='constant_formulas':
            d=out/'cache';d.mkdir()
            subprocess.run(['libreoffice',f'-env:UserInstallation={(out/"cache-profile").as_uri()}','--headless','--convert-to','xlsx','--outdir',str(d),str(fresh/'commerce_review.xlsx')],check=True,capture_output=True)
            v=load_workbook(d/'commerce_review.xlsx',data_only=True)
            for s in w:
                for row in s:
                    for c in row:
                        if c.data_type=='f':c.value='='+str(v[s.title][c.coordinate].value)
        w.save(f)
    scores={}
    for name,folder in cases.items():
        d=out/'judges'/name;d.mkdir(parents=True)
        with (d/'log.txt').open('w') as log:r=subprocess.run([sys.executable,str(root/'tests/judge.py'),'--artifact-dir',str(folder),'--result',str(d/'result.json'),'--reward-file',str(d/'reward.txt')],stdout=log,stderr=subprocess.STDOUT,timeout=360)
        assert r.returncode in (0,1),(name,r.returncode)
        scores[name]=json.loads((d/'result.json').read_text())['reward'];print(name,scores[name],flush=True)
    assert scores['fresh']==scores['golden']==1
    assert all(scores[x]==0 for x in ('empty','blank','corrupt','unchanged'))
    assert all(0<scores[x]<1 for x in ('full_month','wrong_denominator','average_roi','invent_unknown','overwrite_source'))
    assert scores['constant_formulas']<.8
    repeated=[]
    for i in range(3):
        d=out/f'repeat-{i}';d.mkdir()
        subprocess.run([sys.executable,str(root/'tests/judge.py'),'--artifact-dir',str(fresh),'--result',str(d/'result.json'),'--reward-file',str(d/'reward.txt')],check=True,capture_output=True,timeout=360)
        repeated.append(json.loads((d/'result.json').read_text())['reward'])
    assert repeated==[1]*3
    (out/'validation.json').write_text(json.dumps({'task':'workspace-commerce-metric-reconciliation','scores':scores,'repeated_rewards':repeated,'all_assertions_passed':True},ensure_ascii=False,indent=2))
if __name__=='__main__':main()

