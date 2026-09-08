"""仅本题：历史窗口重建、单品净化与来源证据。"""
import argparse
from pathlib import Path
import xlrd
from openpyxl import Workbook,load_workbook
from openpyxl.chart import BarChart,Reference
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter

def main():
    p=argparse.ArgumentParser();p.add_argument('--input-dir',type=Path,required=True);p.add_argument('--output-dir',type=Path,required=True);a=p.parse_args()
    w=Workbook();w.remove(w.active)
    def sheet(name,header):
        s=w.create_sheet(name);s.append(header);return s
    src=load_workbook(next(a.input_dir.glob('*.xlsx')),data_only=True)['tmall 周报']
    s=sheet('周报底稿',['日','销售','PC销售','移动销售','UV','PCUV','移动UV','源转化率','直通车UV','直通车成本','直通车比例','钻展UV','钻展成本','钻展比例','来源'])
    def roi(v):return float(v.split(':')[-1]) if isinstance(v,str) else float(v)
    for col in range(5,21):
        values=[src.cell(r,col).value for r in [3,4,5,6,7,8,9,13,15,16,17,19,20]]
        values[9]=roi(values[9]);values[12]=roi(values[12])
        s.append([src.cell(2,col).value]+values+[f'tmall 周报!{get_column_letter(col)}3:{get_column_letter(col)}20'])
    product=xlrd.open_workbook(str(a.input_dir/'_单品数据_单品数据跟踪.xls'));ps=product.sheet_by_name('醇香流量跟踪')
    s=sheet('单品底稿',['日期','访客','买家','金额','假访客','刷单人数','刷单金额','源实际访客','源实际金额','源实际转化率','源实际买家','源实际客单价','原备注','源行号'])
    for r in range(5,12):
        s.append([xlrd.xldate_as_datetime(ps.cell_value(r,0),product.datemode).date().isoformat()]+[ps.cell_value(r,c) for c in [2,4,5,7,9,10,8,11,12,13,14,40]]+[r+1])
    s=sheet('来源登记',['文件','工作表','期间','用途','边界'])
    for row in [
        ['_店铺数据_tsl项目周报表-7月第三周.xlsx','tmall 周报','第1—7日；另留1—16日','天猫日序指标','年份未确认；不混入京东和11月单品'],
        ['_单品数据_单品数据跟踪.xls','醇香流量跟踪','2015-11-01—07','单品净化','只以工作表名标识，不能由备注确认商品身份'],
        ['_单品数据_单品数据跟踪.xls','碧螺春','同期空数据','身份对照','空白不能补0或冒认另表指标'],
        ['_单品数据_单品数据跟踪.xls','无线流量/PC流量','2015年7月','排除不同窗口','不并入11月指标'],
        ['_工作计划_月_周工作总结计划表.xls','周报/月报','12月26日-01月01日；含2011历史日期','历史流程参考','不能证明当前活动已完成'],
        ['_数据分析_淘宝网店日常统计数表格.xls','Sheet1','未填期间','指标模板参考','目标、精确订单、周去重人数未知']]:s.append(row)
    s=sheet('周报复核',['日','销售','移动占比','UV差异','直通车每UV成本','钻展每UV成本','直通车隐含额','钻展隐含额','估计订单','销售差异'])
    for day in range(1,8):
        r=day+1
        def v(col):return f"SUMIF('周报底稿'!A2:A17,A{r},'周报底稿'!{col}2:{col}17)"
        s.append([day,'='+v('B'),'='+v('D')+'/'+v('B'),'='+v('F')+'+'+v('G')+'-'+v('E'),'='+v('J')+'/'+v('I'),'='+v('M')+'/'+v('L'),'='+v('J')+'*'+v('K'),'='+v('M')+'*'+v('N'),'='+v('E')+'*'+v('H'),'='+v('B')+'-'+v('C')+'-'+v('D')])
    s=sheet('单品净化',['日期','净访客','净买家','净额','净转化率','净客单价','原分母转化率','转化率差','访客残差','金额残差','买家残差'])
    for r in range(2,9):
        def v(col):return f"SUMIF('单品底稿'!A2:A8,A{r},'单品底稿'!{col}2:{col}8)"
        s.append([w['单品底稿'].cell(r,1).value,'='+v('B')+'-'+v('E'),'='+v('C')+'-'+v('F'),'='+v('D')+'-'+v('G'),f'=C{r}/B{r}',f'=D{r}/C{r}','=C'+str(r)+'/'+v('B'),f'=E{r}-'+v('J'),f'=B{r}-'+v('H'),f'=D{r}-'+v('I'),f'=C{r}-'+v('K')])
    s=sheet('经营指标',['指标','结果','口径'])
    def first(col):return f"SUMIF('周报底稿'!A2:A17,\"<=7\",'周报底稿'!{col}2:{col}17)"
    metrics=[
        ('第一周销售',"SUM('周报复核'!B2:B8)",'第1—7日'),
        ('第一周移动销售占比',first('D')+'/B2','金额加权'),
        ('逐日全店UV之和',first('E'),'不是周去重'),
        ('订单估计之和',"SUM('周报复核'!I2:I8)",'四舍五入比例隐含值'),
        ('UV加权转化率估计','B5/B4','估计订单/逐日UV和'),
        ('直通车花费',first('J'),'第1—7日'),
        ('钻展花费',first('M'),'第1—7日'),
        ('直通车每UV成本','B7/'+first('I'),'不称CPC'),
        ('钻展每UV成本','B8/'+first('L'),'不称CPC'),
        ('直通车隐含ROAS估计',"SUM('周报复核'!G2:G8)/B7",'源比例×日成本后汇总，不是利润'),
        ('钻展隐含ROAS估计',"SUM('周报复核'!H2:H8)/B8",'源比例已舍入'),
        ('单品净额',"SUM('单品净化'!D2:D8)",'原额-刷单额'),
        ('单品刷单金额占比',"SUM('单品底稿'!G2:G8)/SUM('单品底稿'!D2:D8)",'总量之比'),
        ('单品净访客口径总转化率',"SUM('单品净化'!C2:C8)/SUM('单品净化'!B2:B8)",'净买家/净访客'),
        ('单品净客单价',"SUM('单品净化'!D2:D8)/SUM('单品净化'!C2:C8)",'净额/净买家'),
        ('第1—16日销售',"SUM('周报底稿'!B2:B17)",'不是第一周'),
        ('第一周占16日销售比','B2/B17','仅同表日序'),
        ('精确订单数',None,'缺失，不能将估计取整冒充真值'),
        ('本周销售目标',None,'模板未提供'),
        ('周去重访客数',None,'跨天身份未知'),
        ('单品净访客之和',"SUM('单品净化'!B2:B8)",'日合计'),
        ('单品净买家之和',"SUM('单品净化'!C2:C8)",'日合计'),
        ('单品原访客口径总转化率',"SUM('单品净化'!C2:C8)/SUM('单品底稿'!B2:B8)",'保留分母差异'),
        ('两种单品总转化率之差','B15-B24','净访客口径减原访客口径')]
    for label,formula,note in metrics:s.append([label,'='+formula if formula else None,note])
    plan=xlrd.open_workbook(str(a.input_dir/'_工作计划_月_周工作总结计划表.xls')).sheet_by_name('周报')
    s=sheet('工作证据',['源行号','事项','解决方法原文','复用边界'])
    for r in [5,6,7,8,11]:s.append([r+1,plan.cell_value(r,2),plan.cell_value(r,8),'仅历史流程参考，不能据此断言2015年11月活动已完成'])
    s=sheet('审计说明',['主题','证据与结论'])
    for row in [
        ['范围','周报第1—7日与D列16日累计不同；单品2015年11月，不同期间不能合并。来源登记保留排除项。'],
        ['分母','醇香流量跟踪M列源实际转化率是净买家/原访客；净访客口径另算，不能覆盖源数。'],
        ['身份','醇香流量跟踪备注碧螺春B淘抢购，而碧螺春工作表同期为空，商品身份不能确认。'],
        ['限制','跨天UV不是周去重；PC加移动不一定可加；ROI仅隐含归因估计，不是利润；订单由已舍入转化率估计，不是精确值。'],
        ['未知','日常统计Sheet1模板未提供目标；工作计划含2011历史活动，是否完成未知，不能编造当期事实。']]:s.append(row)
    chart=BarChart();chart.title='单品原额与净额';chart.add_data(Reference(w['单品底稿'],min_col=4,min_row=1,max_row=8),titles_from_data=True);chart.add_data(Reference(w['单品净化'],min_col=4,min_row=1,max_row=8),titles_from_data=True);chart.set_categories(Reference(w['单品净化'],min_col=1,min_row=2,max_row=8));w['审计说明'].add_chart(chart,'A10')
    for s in w:
        s.freeze_panes='B2';s.auto_filter.ref=s.dimensions;s.sheet_view.showGridLines=False
        s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0;s.sheet_properties.pageSetUpPr.fitToPage=True
        s.print_title_rows='1:1';s.print_options.horizontalCentered=True
        for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='24556A');c.alignment=Alignment(wrap_text=True)
        s.row_dimensions[1].height=35
        for row in s.iter_rows(min_row=2):
            for c in row:
                c.alignment=Alignment(vertical='center',wrap_text=True)
                if isinstance(c.value,(int,float)) or c.data_type=='f':c.number_format='#,##0.00;[Red]-#,##0.00'
            s.row_dimensions[row[0].row].height=30
        for col in range(1,s.max_column+1):s.column_dimensions[get_column_letter(col)].width=19
        s.column_dimensions['A'].width=24
    for title,cols in [('周报底稿',['H','K','N']),('单品底稿',['J']),('周报复核',['C']),('单品净化',['E','G','H'])]:
        for col in cols:
            for c in w[title][col][1:]:c.number_format='0.0000%'
    for r in [3,6,11,12,14,15,18,24,25]:w['经营指标'][f'B{r}'].number_format='0.0000%'
    w['经营指标'].column_dimensions['A'].width=38;w['经营指标'].column_dimensions['C'].width=58
    for r in range(2,26):w['经营指标'].row_dimensions[r].height=24
    for c in 'ABCDE':w['来源登记'].column_dimensions[c].width=38
    for r in range(2,8):w['来源登记'].row_dimensions[r].height=54
    for c in 'BCD':w['工作证据'].column_dimensions[c].width=60
    for r in range(2,7):w['工作证据'].row_dimensions[r].height=60
    w['审计说明'].column_dimensions['B'].width=120
    for r in range(2,7):w['审计说明'].row_dimensions[r].height=56
    w['审计说明'].print_area='A1:F26';w['周报底稿'].column_dimensions['O'].width=36;w['单品底稿'].column_dimensions['M'].width=30
    a.output_dir.mkdir(parents=True,exist_ok=True);w.save(a.output_dir/'commerce_review.xlsx')
if __name__=='__main__':main()
