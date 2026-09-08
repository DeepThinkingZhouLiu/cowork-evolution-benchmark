"""网店题独立阅卷，不读取Golden；核验来源、分母与重算行为。"""
import argparse,json,math,subprocess,tempfile,zipfile
from pathlib import Path
import xlrd
from openpyxl import load_workbook

def equal(x,y):
    if y in (None,''):return x in (None,'')
    if isinstance(y,(int,float)):return isinstance(x,(int,float)) and not isinstance(x,bool) and math.isclose(x,y,rel_tol=1e-8,abs_tol=1e-5)
    return x==y
def calc(path,temp,label):
    dest=temp/label;dest.mkdir()
    r=subprocess.run(['libreoffice',f'-env:UserInstallation={(temp/(label+"-profile")).as_uri()}','--headless','--convert-to','xlsx','--outdir',str(dest),str(path)],capture_output=True,timeout=180)
    if r.returncode or not (dest/path.name).exists():raise RuntimeError('LibreOffice重算失败')
    return load_workbook(dest/path.name,data_only=True)

def main():
    p=argparse.ArgumentParser();p.add_argument('--artifact-dir',type=Path,required=True);p.add_argument('--result',type=Path,required=True);p.add_argument('--reward-file',type=Path,required=True);a=p.parse_args()
    tests=Path(__file__).parent;inp=tests.parent/'data/input_files';rub=json.loads((tests/'rubric.json').read_text());path=a.artifact_dir/'commerce_review.xlsx'
    scores={c['id']:0. for c in rub['criteria']};evidence={}
    def mark(k,checks,text):scores[k]=sum(map(bool,checks))/len(checks) if checks else 0.;evidence[k]=text
    def cells(book,title,rows):
        if title not in book:return [False]
        s=book[title];res=[s.max_row==len(rows)+1]
        for r,row in enumerate(rows,2):res.extend(equal(s.cell(r,c).value,x) for c,x in enumerate(row,1))
        return res
    w=None
    try:
        if path.exists():w=load_workbook(path)
    except (ValueError,KeyError,OSError,zipfile.BadZipFile):pass
    if w:
        mark('R001',['周报底稿' in w and w['周报底稿'].max_row>1],'非空业务底稿与可解析原生XLSX')
        s=load_workbook(next(inp.glob('*.xlsx')),data_only=True)['tmall 周报'];weekly=[]
        for c in range(5,21):
            values=[]
            for r in (2,3,4,5,6,7,8,9,13,15,16,17,19,20):
                x=s.cell(r,c).value
                if r in (16,20):x=float(str(x).removeprefix('1:'))
                values.append(x)
            from openpyxl.utils import get_column_letter
            weekly.append(values+[f'tmall 周报!{get_column_letter(c)}3:{get_column_letter(c)}20'])
        book=xlrd.open_workbook(str(inp/'_单品数据_单品数据跟踪.xls'));s=book.sheet_by_name('醇香流量跟踪');product=[]
        for r in range(5,12):
            date=xlrd.xldate_as_datetime(s.cell_value(r,0),book.datemode).strftime('%Y-%m-%d')
            product.append([date]+[s.cell_value(r,c) for c in (2,4,5,7,9,10,8,11,12,13,14,40)]+[r+1])
        mark('R002',cells(w,'周报底稿',weekly),'16日天猫，混合类型ROI解析，原数与定位')
        mark('R003',cells(w,'单品底稿',product),'七日源值、分母、备注及日期；不冒认空白碧螺春表')
        def derive(week,prod):
            seven=sorted((r for r in week if r[0]<=7),key=lambda r:r[0]);prod=sorted(prod,key=lambda r:r[0])
            wd=[[r[0],r[1],r[3]/r[1],r[5]+r[6]-r[4],r[9]/r[8],r[12]/r[11],r[9]*r[10],r[12]*r[13],r[4]*r[7],r[1]-r[2]-r[3]] for r in seven]
            clean=[]
            for r in prod:
                visitors=r[1]-r[4];buyers=r[2]-r[5];paid=r[3]-r[6]
                clean.append([r[0],visitors,buyers,paid,buyers/visitors,paid/buyers,buyers/r[1],buyers/visitors-r[9],visitors-r[7],paid-r[8],buyers-r[10]])
            total=lambda rows,k:sum(r[k] for r in rows)
            sales=total(seven,1);uv=total(seven,4);cost1=total(seven,9);cost2=total(seven,12);netuv=total(clean,1);netbuyers=total(clean,2);netpay=total(clean,3)
            metrics=[sales,total(seven,3)/sales,uv,total(wd,8),total(wd,8)/uv,cost1,cost2,cost1/total(seven,8),cost2/total(seven,11),total(wd,6)/cost1,total(wd,7)/cost2,netpay,total(prod,6)/total(prod,3),netbuyers/netuv,netpay/netbuyers,total(week,1),sales/total(week,1),None,None,None,netuv,netbuyers,netbuyers/total(prod,1),netbuyers/netuv-netbuyers/total(prod,1)]
            return wd,clean,metrics
        wd,clean,metrics=derive(weekly,product)
        with tempfile.TemporaryDirectory(prefix='commerce-judge-') as td:
            temp=Path(td)
            try:v=calc(path,temp,'base')
            except (RuntimeError,subprocess.TimeoutExpired) as e:
                a.result.parent.mkdir(parents=True,exist_ok=True);a.result.write_text(json.dumps({'status':'infrastructure_failure','error':str(e)},ensure_ascii=False));raise SystemExit(2)
            mark('R004',cells(v,'周报复核',wd),'按日UV成本、隐含额、估计订单与非可加差异')
            mark('R005',cells(v,'单品净化',clean),'两种访客分母、净额净人数与源残差')
            mark('R006',[equal(v['经营指标'].cell(r,2).value,x) if '经营指标' in v else False for r,x in enumerate(metrics,2)],'24总控，金额加权比率，未知保持空白')
            coords=[('周报复核',f'{c}{r}') for r in range(2,9) for c in 'BCDEFGHIJ']+[('单品净化',f'{c}{r}') for r in range(2,9) for c in 'BCDEFGHIJK']+[('经营指标',f'B{r}') for r in range(2,26) if r not in (19,20,21)]
            mark('R007',[s in w and w[s][c].data_type=='f' for s,c in coords],'计算项必须原生公式，未知不强制公式')
            changed=[]
            required=['周报底稿','单品底稿','来源登记','周报复核','单品净化','经营指标','工作证据','审计说明']
            if all(s in w for s in required) and w['周报底稿']['A2'].value==1 and w['单品底稿']['A2'].value=='2015-11-01':
                q=load_workbook(path);wk=[r[:] for r in weekly];pr=[r[:] for r in product]
                for col,index,delta in [('B',1,1000),('C',2,1000),('J',9,100)]:
                    wk[0][index]+=delta;q['周报底稿'][col+'2']=wk[0][index]
                for col,index,delta in [('E',4,100),('G',6,50)]:
                    pr[0][index]+=delta;q['单品底稿'][col+'2']=pr[0][index]
                for name,cols in [('周报底稿',15),('单品底稿',14)]:
                    s=q[name];first=[s.cell(2,c).value for c in range(1,cols+1)];last=[s.cell(8,c).value for c in range(1,cols+1)]
                    for c in range(1,cols+1):s.cell(2,c).value=last[c-1];s.cell(8,c).value=first[c-1]
                probe=temp/'probe.xlsx';q.save(probe)
                try:pv=calc(probe,temp,'changed')
                except (RuntimeError,subprocess.TimeoutExpired) as e:
                    a.result.parent.mkdir(parents=True,exist_ok=True);a.result.write_text(json.dumps({'status':'infrastructure_failure','error':str(e)},ensure_ascii=False));raise SystemExit(2)
                nw,nc,nm=derive(wk,pr)
                for name,before,after in [('周报复核',wd,nw),('单品净化',clean,nc)]:
                    for r,(old,new) in enumerate(zip(before,after),2):
                        for c,(b,x) in enumerate(zip(old,new),1):
                            if isinstance(x,(int,float)) and not equal(b,x):changed.append(equal(pv[name].cell(r,c).value,x))
                changed.extend(equal(pv['经营指标'].cell(r,2).value,x) for r,(b,x) in enumerate(zip(metrics,nm),2) if not equal(b,x))
            mark('R008',changed,'销售/广告成本/假访客/刷单额变化与换行后，重算受影响项')
        s=xlrd.open_workbook(str(inp/'_工作计划_月_周工作总结计划表.xls')).sheet_by_name('周报')
        checks=[]
        if '工作证据' in w:
            for row,r in enumerate([5,6,7,8,11],2):
                checks.extend([w['工作证据'].cell(row,1).value==r+1,w['工作证据'].cell(row,2).value==s.cell_value(r,2),w['工作证据'].cell(row,3).value==s.cell_value(r,8)])
                text=str(w['工作证据'].cell(row,4).value or '');checks.append(('不能' in text or '不代表' in text) and '历史' in text)
        registry=[' '.join(str(c or '') for c in row) for row in w['来源登记'].iter_rows(min_row=2,values_only=True)] if '来源登记' in w else []
        for tokens in [('tmall 周报','1','7'),('醇香流量跟踪','2015'),('碧螺春','空'),('无线','PC','7'),('工作计划','历史'),('日常统计','模板')]:
            checks.append(any(all(t in line for t in tokens) for line in registry))
        mark('R009',checks,'历史原文定位与六类期间登记，不能冒充执行完成')
        lines=[' '.join(str(x or '') for x in row) for row in w['审计说明'].iter_rows(min_row=2,values_only=True)] if '审计说明' in w else []
        checks=[]
        for tokens in [('16','7','期间'),('净买家','原访客','源'),('醇香','碧螺春','身份'),('UV','估计','利润'),('目标','历史','未知')]:
            checks.append(any(all(t in line for t in tokens) and len(line)>30 for line in lines))
        checks.extend([any(s._charts for s in w),not w._external_links,all(s in w and w[s].freeze_panes and w[s].auto_filter.ref for s in required)])
        mark('R010',checks,'关键口径边界的有限关键词核验与图表可编辑性；不评审美')
    score=round(sum(c['weight']*scores[c['id']] for c in rub['criteria']),6)
    if scores['R001']<1:score=0
    result={'task_id':rub['task_id'],'judge_version':'commerce-metric-reconciliation-v1','rubric_version':rub['rubric_version'],'reward':score,'passed':score>=rub['pass_threshold'],'criterion_results':[{'criterion_id':k,'score':round(v,6),'method':'deterministic','evidence':evidence.get(k,'未满足')} for k,v in scores.items()]}
    a.result.parent.mkdir(parents=True,exist_ok=True);a.result.write_text(json.dumps(result,ensure_ascii=False,indent=2));a.reward_file.parent.mkdir(parents=True,exist_ok=True);a.reward_file.write_text(str(score)+'\n')
    raise SystemExit(0 if result['passed'] else 1)
if __name__=='__main__':main()

