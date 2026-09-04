#!/usr/bin/env python3
import argparse, re, shutil
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

def parse_md(path):
    t = path.read_text(encoding='utf-8')
    n = int(re.search(r'_(\d+)\.md$', path.name).group(1))
    def field(label):
        m = re.search(rf'\*\*{re.escape(label)}\*\*:\s*(.+)', t)
        return m.group(1).strip() if m else ''
    lead = re.search(r'Lead generation \((\d+) target\)', t, re.I)
    manager = re.search(r'- Event Manager:\s*(.+)', t)
    marketing = re.search(r'- Marketing Lead:\s*(.+)', t)
    status = re.search(r'\*\*Status\*\*:\s*(.+)', t)
    date = datetime.strptime(field('Date'), '%Y-%m-%d')
    return [path.name, f'Event {n}', field('Event Name'), date, field('Location'), int(field('Expected Attendees')), int(lead.group(1)), manager.group(1).strip(), marketing.group(1).strip(), status.group(1).strip()]

def style(ws):
    ws.freeze_panes = 'A2'; ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = f'A1:J{ws.max_row}'
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78'); c.alignment = Alignment(horizontal='center')
    widths = [22, 12, 28, 14, 20, 20, 23, 20, 24, 18]
    for i,w in enumerate(widths,1): ws.column_dimensions[chr(64+i)].width = w
    tab = Table(displayName='EventPlanTable', ref=f'A1:J{ws.max_row}'); tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True); ws.add_table(tab)

def build(input_dir, output_dir):
    output_dir.mkdir(parents=True, exist_ok=True)
    rows = [parse_md(p) for p in input_dir.glob('event_plan_*.md')]
    rows.sort(key=lambda r: int(r[1].split()[-1]))
    if len(rows) != 60: raise ValueError(f'expected 60 markdown files, got {len(rows)}')
    wb = load_workbook(input_dir/'event_plan_v1.xlsx')
    ws = wb.active; ws.title = 'Event Plan'
    ws.delete_rows(1, ws.max_row)
    ws.append(['Filename','Event ID','Event Name','Date','Location','Expected Attendees','Lead Generation Target','Event Manager','Marketing Lead','Status'])
    for row in rows:
        ws.append(row); ws.cell(ws.max_row,4).number_format='yyyy-mm-dd'; ws.cell(ws.max_row,6).number_format='#,##0'; ws.cell(ws.max_row,7).number_format='#,##0'
    style(ws)
    if 'Source Notes' in wb.sheetnames: del wb['Source Notes']
    notes = wb.create_sheet('Source Notes'); notes.append(['Field','Value']); notes.append(['Source','Workspace-Bench-Lite task 159; 60 event_plan_*.md files']); notes.append(['Method','Rebuilt from each Markdown; v1 used only as editable workbook shell']); notes.append(['Status','pending_human_review / uncalibrated'])
    for c in notes[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F4E78')
    notes.column_dimensions['A'].width=22; notes.column_dimensions['B'].width=100; notes.freeze_panes='A2'
    out = output_dir/'event_plan_v2.xlsx'; wb.save(out); return out

if __name__ == '__main__':
    ap=argparse.ArgumentParser(); ap.add_argument('--input-dir',type=Path,required=True); ap.add_argument('--output-dir',type=Path,required=True); a=ap.parse_args(); build(a.input_dir,a.output_dir)
