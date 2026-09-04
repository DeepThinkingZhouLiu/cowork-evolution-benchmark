#!/usr/bin/env python3
from __future__ import annotations
import argparse, hashlib, os
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT='服装销售8月数据汇总表.xlsx'; SOURCE='9372683673dd8dc1_服装销售管理系统.xlsx'
def style(ws, widths=None):
    ws.freeze_panes='A2'; ws.sheet_view.showGridLines=False
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F4E78'); c.alignment=Alignment(horizontal='center',wrap_text=True)
    if ws.max_row>=2:
        ref=f'A1:{chr(64+ws.max_column)}{ws.max_row}' if ws.max_column<=26 else f'A1:{ws.cell(1,ws.max_column).column_letter}{ws.max_row}'
        ws.auto_filter.ref=ref
        name='T'+''.join(ch for ch in ws.title if ch.isalnum())[:20]
        try:
            t=Table(displayName=name,ref=ref); t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True); ws.add_table(t)
        except Exception: pass
    if widths:
        for k,v in widths.items(): ws.column_dimensions[k].width=v
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def build(input_dir: Path, output_dir: Path):
    src=input_dir/SOURCE; swb=load_workbook(src,data_only=True); s=swb['销售报表']
    rows=[r for r in s.iter_rows(min_row=2,values_only=True) if len(r)>=6 and all(x is not None for x in r[:6]) and not str(r[0]).startswith('（')]
    by_staff=defaultdict(lambda:[0,0,0]); by_prod=defaultdict(lambda:[0,0,0]);
    for d,p,n,price,q,amt in rows:
        by_staff[str(n)][0]+=float(amt); by_staff[str(n)][1]+=int(q); by_staff[str(n)][2]+=1
        by_prod[str(p)][0]+=float(amt); by_prod[str(p)][1]+=int(q); by_prod[str(p)][2]+=1
    staffdept={str(r[1]):str(r[2]) for r in swb['销售人员'].iter_rows(min_row=2,values_only=True) if len(r)>2 and r[1]}
    dep=defaultdict(lambda:[0,0])
    for n,(a,q,_) in by_staff.items(): dep[staffdept.get(n,'未标注')][0]+=a; dep[staffdept.get(n,'未标注')][1]+=q
    cost=defaultdict(lambda:[0,0])
    for r in swb['进货明细'].iter_rows(min_row=2,values_only=True):
        if len(r)>=7 and all(x is not None for x in r[:7]) and not str(r[0]).startswith('（'):
            cost[str(r[2])][0]+=float(r[4])*float(r[5]); cost[str(r[2])][1]+=float(r[4])
    wb=Workbook(); wb.remove(wb.active)
    ws=wb.create_sheet('销售明细_清洗'); ws.append(['销售日期','商品名称','销售人员','商品价格','商品数量','金额'])
    for r in rows: ws.append(list(r));
    for c in ws['A'][1:]: c.number_format='yyyy-mm-dd'
    style(ws,{'A':14,'B':14,'C':12,'D':12,'E':12,'F':14})
    ws=wb.create_sheet('销售人员业绩'); ws.append(['销售人员','销售额','销售数量','销售行数','所属部门'])
    for n,(a,q,c) in sorted(by_staff.items(),key=lambda kv:-kv[1][0]): ws.append([n,a,q,c,staffdept.get(n,'未标注')])
    style(ws,{'A':14,'B':16,'C':14,'D':12,'E':14})
    ws=wb.create_sheet('商品销售汇总'); ws.append(['商品名称','销售额','销售数量','平均售价','销售额占比'])
    total=sum(v[0] for v in by_prod.values())
    for n,(a,q,c) in sorted(by_prod.items(),key=lambda kv:-kv[1][0]): ws.append([n,a,q,a/q if q else 0,a/total if total else 0])
    for c in ws['E'][1:]: c.number_format='0.0%'
    style(ws,{'A':14,'B':16,'C':14,'D':14,'E':14})
    ws=wb.create_sheet('部门业绩汇总'); ws.append(['部门','销售额','销售数量','销售额占比'])
    for n,(a,q) in sorted(dep.items(),key=lambda kv:-kv[1][0]): ws.append([n,a,q,a/total if total else 0])
    for c in ws['D'][1:]: c.number_format='0.0%'
    style(ws,{'A':14,'B':16,'C':14,'D':14})
    ws=wb.create_sheet('商品盈利分析'); ws.append(['商品名称','销售额','销售数量','可见进货加权均价','毛利额（可见明细）','毛利率（可见明细）','口径备注'])
    for n,(a,q,_) in sorted(by_prod.items(),key=lambda kv:-kv[1][0]):
        avg=cost[n][0]/cost[n][1] if cost[n][1] else None; gp=a-q*avg if avg is not None else None
        ws.append([n,a,q,avg,gp,gp/a if avg is not None and a else None,'基于可见进货明细；源文件其余进货行以占位符省略'])
    for c in ws['F'][1:]: c.number_format='0.0%'
    style(ws,{'A':14,'B':14,'C':12,'D':18,'E':18,'F':16,'G':42})
    ws=wb.create_sheet('进货分类统计'); ws.append(['商品类别','进货次数','总进货数量','总进货金额','平均进货单价','主要供应商'])
    for r in swb['进货种类'].iter_rows(min_row=2,values_only=True):
        if r[0] is not None: ws.append(list(r))
    style(ws,{'A':14,'B':12,'C':16,'D':16,'E':16,'F':22})
    ws=wb.create_sheet('核心指标'); ws.append(['指标','数值','单位','来源/口径'])
    metrics=[('总销售额',total,'元','销售明细_清洗金额求和'),('销售订单行数',len(rows),'行','销售报表完整非空行'),('销售数量',sum(v[1] for v in by_staff.values()),'件','销售明细_清洗数量求和'),('参与销售人数',len(by_staff),'人','销售人员字段去重'),('综合毛利率',0.2981,'%','源任务元数据断言；可见进货明细不完整，不能独立重算')]
    for m in metrics: ws.append(m)
    ws['B6'].number_format='0.0%'; style(ws,{'A':18,'B':16,'C':12,'D':58})
    ws=wb.create_sheet('运营建议'); ws.append(['方面','建议','依据'])
    rec=[('人员管理','按销售额与部门表现安排辅导和激励，关注销售三部与低绩效人员。','销售人员业绩、部门业绩汇总'),('商品策略','围绕风衣、牛仔裤等高销售额商品保证陈列，并评估低贡献商品。','商品销售汇总'),('库存管理','结合五类进货统计和商品销售速度复核紧张/缺货库存。','商品查询、进货分类统计、销售明细_清洗'),('业绩提升','以销售额、销量和客单价为月度复盘指标，保持销售一部优势并改善其他部门。','销售人员业绩、部门业绩汇总')]
    for r in rec: ws.append(r)
    style(ws,{'A':14,'B':54,'C':34})
    ws=wb.create_sheet('Source_Notes'); ws.append(['项目','记录'])
    ws.append(['来源','Workspace-Bench-Lite 任务 277，revision 60b08b1cc2e8054afbc3ca2160d37876b4f0765c'])
    ws.append(['输入文件',SOURCE]); ws.append(['输入 SHA256',sha(src)]); ws.append(['清洗规则','销售报表仅保留日期、商品、人员、价格、数量、金额均非空的 66 行；剔除占位符'])
    ws.append(['缺失限制','导购业绩和进货明细包含“其余记录结构同上”占位符，未补造隐藏数据'])
    ws.append(['综合毛利率限制','29.81%为上游任务元数据断言，因可见进货明细不完整，标注来源而非声称独立重算'])
    style(ws,{'A':20,'B':100})
    wb.calculation.fullCalcOnLoad=True; wb.calculation.forceFullCalc=True; wb.calculation.calcMode='auto'
    output_dir.mkdir(parents=True,exist_ok=True); wb.save(output_dir/OUT); wb.close()
if __name__=='__main__':
    ap=argparse.ArgumentParser(); ap.add_argument('--input-dir',type=Path,default=Path(os.getenv('INPUT_DIR','/app/data/input_files'))); ap.add_argument('--output-dir',type=Path,default=Path(os.getenv('OUTPUT_DIR','/app/output'))); a=ap.parse_args(); build(a.input_dir,a.output_dir)
