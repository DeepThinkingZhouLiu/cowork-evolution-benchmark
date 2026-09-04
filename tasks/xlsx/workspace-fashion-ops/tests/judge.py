#!/usr/bin/env python3
from __future__ import annotations
import argparse, hashlib, json, zipfile
from pathlib import Path
from collections import defaultdict
from openpyxl import load_workbook
OUT='服装销售8月数据汇总表.xlsx'; SRC='9372683673dd8dc1_服装销售管理系统.xlsx'
def crit(i,ok,e=''): return {'criterion_id':i,'score':1.0 if ok else 0.0,'evidence':e,'method':'deterministic'}
def main():
 ap=argparse.ArgumentParser(); ap.add_argument('--artifact-dir',type=Path,required=True); ap.add_argument('--result',type=Path,required=True); ap.add_argument('--reward-file',type=Path,required=True); a=ap.parse_args(); root=Path(__file__).parents[1]; inp=root/'data/input_files'/SRC; out=a.artifact_dir/OUT; rs=[]
 try:
  valid=out.is_file() and zipfile.is_zipfile(out); rs.append(crit('R001',valid,'有效 XLSX' if valid else '缺失/损坏'))
  if not valid: raise ValueError('invalid')
  wi=load_workbook(inp,data_only=True); wo=load_workbook(out,data_only=True); s=wi['销售报表']; rows=[r for r in s.iter_rows(min_row=2,values_only=True) if len(r)>=6 and all(x is not None for x in r[:6]) and not str(r[0]).startswith('（')]
  sales=sum(float(r[5]) for r in rows); qty=sum(int(r[4]) for r in rows); people={str(r[2]) for r in rows}; products=defaultdict(lambda:[0,0])
  for r in rows: products[str(r[1])][0]+=float(r[5]); products[str(r[1])][1]+=int(r[4])
  required=['销售明细_清洗','销售人员业绩','商品销售汇总','部门业绩汇总','商品盈利分析','进货分类统计','核心指标','运营建议','Source_Notes']; names=wo.sheetnames
  sh=wo['销售明细_清洗']; clean= list(sh.iter_rows(min_row=2,values_only=True)); clean=[r for r in clean if r[0] is not None]
  rs.append(crit('R002',all(n in names for n in required) and len(clean)==66 and not any('其余' in str(x) for r in clean for x in r if x is not None),f'清洗行={len(clean)}'))
  cm=wo['核心指标']; vals={str(r[0]):r[1] for r in cm.iter_rows(min_row=2,values_only=True) if r[0]}; totals_ok=abs(float(vals.get('总销售额',0))-1233140)<.01 and int(vals.get('销售订单行数',0))==66 and int(vals.get('销售数量',0))==4285 and int(vals.get('参与销售人数',0))==10
  rs.append(crit('R003',totals_ok,f'指标={vals}'))
  sp=wo['销售人员业绩']; sr=[r for r in sp.iter_rows(min_row=2,values_only=True) if r[0]]; order_ok=len(sr)==10 and all(float(sr[i][1])>=float(sr[i+1][1]) for i in range(len(sr)-1)) and {str(r[0]) for r in sr}==people
  rs.append(crit('R004',order_ok,f'人员={len(sr)}'))
  pp=wo['商品销售汇总']; pr=[r for r in pp.iter_rows(min_row=2,values_only=True) if r[0]]; prod_ok=len(pr)==10 and {str(r[0]) for r in pr}==set(products) and all(abs(float(next(v for k,v in products.items() if k==str(r[0]))[0])-float(r[1]))<.01 for r in pr)
  rs.append(crit('R005',prod_ok,f'商品={len(pr)}'))
  dp=wo['部门业绩汇总']; dr=[r for r in dp.iter_rows(min_row=2,values_only=True) if r[0]]; rs.append(crit('R006',len(dr)==3 and abs(sum(float(r[1]) for r in dr)-sales)<.01,'部门汇总'))
  ic=wo['进货分类统计']; ir=[r for r in ic.iter_rows(min_row=2,values_only=True) if r[0]]; gp=wo['商品盈利分析']; text=' '.join(str(c.value or '') for row in gp.iter_rows() for c in row); rs.append(crit('R007',len(ir)==5 and '可见进货明细' in text,'五类和限制说明'))
  fmt=all(wo[n].freeze_panes=='A2' and bool(wo[n].tables) for n in ['销售明细_清洗','销售人员业绩','商品销售汇总']) and wo['运营建议'].max_row>=5; rs.append(crit('R008',fmt,'表格/冻结/建议'))
  note=' '.join(str(c.value or '') for row in wo['Source_Notes'].iter_rows() for c in row); rs.append(crit('R009','29.81%' in note or '0.2981' in note,'来源限制'))
  reward=sum(r['score']*w for r,w in zip(rs,[.1,.2,.2,.1,.1,.08,.08,.08,.06])); payload={'task_id':'cowork-evo/workspace-fashion-ops','rubric_version':'workspace-fashion-ops-v0.1','criterion_results':rs,'reward':reward,'passed':all(r['score']==1 for r in rs[:5]) and reward>=.8}; a.result.parent.mkdir(parents=True,exist_ok=True); a.result.write_text(json.dumps(payload,ensure_ascii=False,indent=2)); a.reward_file.parent.mkdir(parents=True,exist_ok=True); a.reward_file.write_text(f'{reward:.6f}\n'); raise SystemExit(0 if payload['passed'] else 1)
 except Exception as e:
  payload={'task_id':'cowork-evo/workspace-fashion-ops','criterion_results':rs+[crit('R001',False,str(e))],'reward':0,'passed':False}; a.result.parent.mkdir(parents=True,exist_ok=True); a.result.write_text(json.dumps(payload,ensure_ascii=False,indent=2)); a.reward_file.parent.mkdir(parents=True,exist_ok=True); a.reward_file.write_text('0\n'); raise SystemExit(1)
if __name__=='__main__': main()
