"""HR群体分析专属反例，验证分母、统一权重与公式联动。"""
import argparse,json,subprocess,sys
from pathlib import Path
from openpyxl import Workbook,load_workbook

def main():
    p=argparse.ArgumentParser();p.add_argument('--output',type=Path,required=True);a=p.parse_args()
    out=a.output.resolve();out.mkdir(parents=True,exist_ok=False);root=Path(__file__).resolve().parents[1]
    fresh=out/'fresh'
    subprocess.run([sys.executable,str(root/'solution/solve.py'),'--input-dir',str(root/'data/input_files'),'--output-dir',str(fresh)],check=True)
    cases={'fresh':fresh,'golden':root/'solution/Golden'}
    for name in ['empty','blank','corrupt','unchanged','wrong_denominator','own_weights','swapped_income','wrong_labels','double_count','constant_formulas']:
        d=out/name;d.mkdir();cases[name]=d;f=d/'hr_review.xlsx'
        if name=='empty':continue
        if name=='blank':Workbook().save(f);continue
        if name=='corrupt':f.write_bytes(b'broken');continue
        w=load_workbook(fresh/'hr_review.xlsx')
        if name=='unchanged':
            for title in list(w.sheetnames):
                if title!='联合底稿':del w[title]
        if name=='wrong_denominator':
            for r in range(2,5):w['分组统计'][f'E{r}']=f'=D{r}/1470'
        if name=='own_weights':
            for r in range(2,5):w['职级标准化'][f'J{r}']=f'=I{r}'
        if name=='swapped_income':
            for r in range(2,1472):w['联合底稿'][f'F{r}']=w['联合底稿'][f'G{r}'].value
        if name=='wrong_labels':
            for r in range(2,1472):
                if w['联合底稿'][f'I{r}'].value==3:w['联合底稿'][f'R{r}']='High'
        if name=='double_count':
            for row in list(w['联合底稿'].iter_rows(min_row=2,values_only=True)):w['联合底稿'].append(row)
        if name=='constant_formulas':
            cache=out/'cached';cache.mkdir()
            subprocess.run(['libreoffice',f'-env:UserInstallation={(out/"cache-profile").as_uri()}','--headless','--convert-to','xlsx','--outdir',str(cache),str(fresh/'hr_review.xlsx')],check=True,capture_output=True)
            v=load_workbook(cache/'hr_review.xlsx',data_only=True)
            for s in w:
                for row in s:
                    for c in row:
                        if c.data_type=='f':
                            x=v[s.title][c.coordinate].value;c.value='=""' if x in (None,'') else '='+str(x)
        w.save(f)
    scores={}
    for name,folder in cases.items():
        d=out/'judges'/name;d.mkdir(parents=True)
        with (d/'log.txt').open('w') as log:
            r=subprocess.run([sys.executable,str(root/'tests/judge.py'),'--artifact-dir',str(folder),'--result',str(d/'result.json'),'--reward-file',str(d/'reward.txt')],stdout=log,stderr=subprocess.STDOUT,timeout=360)
        assert r.returncode in (0,1),(name,r.returncode)
        scores[name]=json.loads((d/'result.json').read_text())['reward'];print(name,scores[name],flush=True)
    assert scores['fresh']==scores['golden']==1
    assert all(scores[x]==0 for x in ['empty','blank','corrupt'])
    assert scores['unchanged']<.3
    assert all(0<scores[x]<1 for x in ['wrong_denominator','own_weights','swapped_income','wrong_labels','double_count'])
    assert scores['constant_formulas']<.8
    repeated=[]
    for i in range(3):
        d=out/f'repeat-{i}';d.mkdir()
        subprocess.run([sys.executable,str(root/'tests/judge.py'),'--artifact-dir',str(fresh),'--result',str(d/'result.json'),'--reward-file',str(d/'reward.txt')],check=True,capture_output=True,timeout=360)
        repeated.append(json.loads((d/'result.json').read_text())['reward'])
    assert repeated==[1]*3
    (out/'validation.json').write_text(json.dumps({'task':'workspace-hr-attrition-stratification','scores':scores,'repeated_rewards':repeated,'all_assertions_passed':True},indent=2,ensure_ascii=False))
if __name__=='__main__':main()

