"""HR题参考解：按编号联表，公开样本仅用于群体描述。"""
import argparse,csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.chart import BarChart,Reference
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter as letter
def main():
    p=argparse.ArgumentParser();p.add_argument('--input-dir',type=Path,required=True);p.add_argument('--output-dir',type=Path,required=True);a=p.parse_args()
    data={}
    for f in a.input_dir.glob('*.csv'):
        with f.open(encoding='utf-8-sig') as stream:data[f.name]=list(csv.DictReader(stream))
    master=data['employee_master_list.csv'];index={name:{r['EmployeeNumber']:r for r in rows} for name,rows in data.items()};mi=index['employee_master_list.csv']
    w=Workbook();w.remove(w.active);raw=w.create_sheet('联合底稿')
    raw.append(['EmployeeNumber','Department','JobRole','JobLevel','Attrition','MonthlyIncome','MonthlyRate','OverTime','PerformanceRating','PercentSalaryHike','JobSatisfaction','WorkLifeBalance','YearsAtCompany','DistanceFromHome','EnvironmentSatisfaction','RelationshipSatisfaction','JobInvolvement','PerformanceLabel','JobSatisfactionLabel'])
    perf_labels={1:'Low',2:'Good',3:'Excellent',4:'Outstanding'};sat_labels={1:'Low',2:'Medium',3:'High',4:'Very High'}
    for person in sorted(master,key=lambda r:int(r['EmployeeNumber'])):
        ident=person['EmployeeNumber'];comp=index['compensation_details.csv'][ident];perf=index['performance_review_2024.csv'][ident];sat=index['employee_satisfaction_survey.csv'][ident];commute=index['commute_and_tenure.csv'][ident]
        raw.append([int(ident),person['Department'],person['JobRole'],int(person['JobLevel']),person['Attrition'],int(comp['MonthlyIncome']),int(comp['MonthlyRate']),comp['OverTime'],int(perf['PerformanceRating']),int(perf['PercentSalaryHike']),int(sat['JobSatisfaction']),int(sat['WorkLifeBalance']),int(commute['YearsAtCompany']),int(commute['DistanceFromHome']),int(sat['EnvironmentSatisfaction']),int(sat['RelationshipSatisfaction']),int(sat['JobInvolvement']),perf_labels[int(perf['PerformanceRating'])],sat_labels[int(sat['JobSatisfaction'])]])
    audit=w.create_sheet('主键审计');audit.append(['文件','原行数','唯一编号数','缺失编号','多出编号','同名字段冲突'])
    for name,rows in sorted(data.items()):
        keys=set(index[name]);conflicts=sum(r[k]!=mi[r['EmployeeNumber']][k] for r in rows if r['EmployeeNumber'] in mi for k in r.keys()&mi[r['EmployeeNumber']].keys() if k!='EmployeeNumber')
        audit.append([name,len(rows),len(keys),len(set(mi)-keys),len(keys-set(mi)),conflicts])
    def rg(c):return f"'联合底稿'!{c}$2:{c}$"+str(raw.max_row)
    stats=w.create_sheet('分组统计');stats.append(['维度','组值','人数','离职数','离职比例','收入均值','工作满意均值','绩效均值','低工作满意比例','加班比例'])
    departments=sorted({r['Department'] for r in master});roles=sorted({r['JobRole'] for r in master})
    for dim,groups,c in [('Department',departments,'B'),('JobRole',roles,'C'),('Attrition',['No','Yes'],'E')]:
        for group in groups:
            r=stats.max_row+1;key=f'B{r}'
            stats.append([dim,group,f'=COUNTIF({rg(c)},{key})',f'=COUNTIFS({rg(c)},{key},{rg("E")},"Yes")',f'=D{r}/C{r}',f'=AVERAGEIF({rg(c)},{key},{rg("F")})',f'=AVERAGEIF({rg(c)},{key},{rg("K")})',f'=AVERAGEIF({rg(c)},{key},{rg("I")})',f'=COUNTIFS({rg(c)},{key},{rg("K")},"<=2")/C{r}',f'=COUNTIFS({rg(c)},{key},{rg("H")},"Yes")/C{r}'])
    strata=w.create_sheet('分层统计');strata.append(['部门','职级','加班','人数','离职数','离职比例','收入均值','工作满意均值'])
    for dept in departments:
        for level in range(1,6):
            for overtime in ['No','Yes']:
                r=strata.max_row+1;criteria=f'{rg("B")},A{r},{rg("D")},B{r},{rg("H")},C{r}'
                strata.append([dept,level,overtime,f'=COUNTIFS({criteria})',f'=COUNTIFS({criteria},{rg("E")},"Yes")',f'=IF(D{r}=0,"",E{r}/D{r})',f'=IF(D{r}=0,"",AVERAGEIFS({rg("F")},{criteria}))',f'=IF(D{r}=0,"",AVERAGEIFS({rg("K")},{criteria}))'])
    standard=w.create_sheet('职级标准化');standard.append(['职级','全员人数','共同权重']+departments+[None,'部门','原始比例','标准化比例','差值'])
    for level in range(1,6):
        r=level+1;standard.cell(r,1,level);standard.cell(r,2,f'=COUNTIF({rg("D")},A{r})');standard.cell(r,3,f'=B{r}/SUM(B$2:B$6)')
        for c,dept in enumerate(departments,4):standard.cell(r,c,f'=COUNTIFS({rg("B")},{letter(c)}$1,{rg("D")},A{r},{rg("E")},"Yes")/COUNTIFS({rg("B")},{letter(c)}$1,{rg("D")},A{r})')
    for r,dept in enumerate(departments,2):
        c=letter(r+2);standard.cell(r,8,dept);standard.cell(r,9,f'=COUNTIFS({rg("B")},H{r},{rg("E")},"Yes")/COUNTIF({rg("B")},H{r})');standard.cell(r,10,f'=SUMPRODUCT(C$2:C$6,{c}$2:{c}$6)');standard.cell(r,11,f'=J{r}-I{r}')
    metrics=w.create_sheet('审查指标');metrics.append(['指标','结果'])
    expressions=[('唯一员工数',f'COUNTA({rg("A")})'),('离职数',f'COUNTIF({rg("E")},"Yes")'),('全员离职比例','B3/B2'),('部门比例简单平均',"AVERAGE('分组统计'!E2:E4)"),('部门人数加权比例',"SUMPRODUCT('分组统计'!C2:C4,'分组统计'!E2:E4)/SUM('分组统计'!C2:C4)"),('在职平均收入',f'AVERAGEIF({rg("E")},"No",{rg("F")})'),('离职平均收入',f'AVERAGEIF({rg("E")},"Yes",{rg("F")})'),('在职减离职收入','B7-B8'),('绩效评级至少3人数',f'COUNTIF({rg("I")},">=3")'),('低工作满意人数',f'COUNTIF({rg("K")},"<=2")')]
    for label,f in expressions:metrics.append([label,'='+f])
    notes=w.create_sheet('判读与边界');notes.append(['主题','证据','结论边界','后续核实'])
    for row in [
        ['花名册去重','三部门63+961+446=1470，与主表同一批编号','不得将主表和部门表相加为2940人；部门未覆盖全员不是缺档','主键唯一、并集一致，同名字段冲突0'],
        ['量表含义','绩效3=Excellent；工作满意3=High；加薪11代表11%','绩效高不等于满意度高；MonthlyRate不是MonthlyIncome','依据两份codebook解释，不自行编币种'],
        ['共同权重与构成','部门职级分布不同；共同权重来自全员5个职级','部门比率简单平均不等于人数加权；标准化只是构成调整','同时报告原始/标准化及各层样本数'],
        ['非因果与小样本','群体离职、工资与满意度只是横截面同现','不能断言加薪会减少离职，也不能从小样本给出因果或个人辞退建议','需要时序、政策和额外证据；不使用年龄性别推断风险']]:notes.append(row)
    chart=BarChart();chart.title='原始与共同职级标准化离职比例';chart.add_data(Reference(standard,min_col=9,max_col=10,min_row=1,max_row=4),titles_from_data=True);chart.set_categories(Reference(standard,min_col=8,min_row=2,max_row=4));chart.width=24;chart.height=11;standard.add_chart(chart,'A10')
    for s in w:
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
        for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='24476B');c.alignment=Alignment(wrap_text=True,vertical='center')
        s.row_dimensions[1].height=34
        for row in s.iter_rows(min_row=2):
            for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
        for c in range(1,s.max_column+1):s.column_dimensions[letter(c)].width=21
        if s.title=='联合底稿':
            for c in 'BC':s.column_dimensions[c].width=32
        if s.title=='判读与边界':
            for c in 'BCD':s.column_dimensions[c].width=52
            for r in range(2,6):s.row_dimensions[r].height=70
        s.sheet_properties.pageSetUpPr.fitToPage=True;s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0;s.print_title_rows='1:1'
    for s,cols in [(stats,'EIJ'),(strata,'F')]:
        for c in cols:
            for r in range(2,s.max_row+1):s[f'{c}{r}'].number_format='0.00%'
    for c in 'CDEF':
        for r in range(2,7):standard[f'{c}{r}'].number_format='0.00%'
    for c in 'IJK':
        for r in range(2,5):standard[f'{c}{r}'].number_format='0.00%'
    standard.column_dimensions['E'].width=32;standard.column_dimensions['H'].width=32
    for s,cols in [(stats,'FGH'),(strata,'GH')]:
        for c in cols:
            for r in range(2,s.max_row+1):s[f'{c}{r}'].number_format='#,##0.00'
    for r in [4,5,6]:metrics[f'B{r}'].number_format='0.00%'
    for r in [7,8,9]:metrics[f'B{r}'].number_format='#,##0.00'
    a.output_dir.mkdir(parents=True,exist_ok=True);w.save(a.output_dir/'hr_review.xlsx')
if __name__=='__main__':main()
