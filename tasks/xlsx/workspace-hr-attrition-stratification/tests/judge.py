"""HR题独立阅卷：八CSV主键审计、群体分母、职级直接标准化与变更检查。"""
import argparse,csv,json,math,statistics,subprocess,tempfile,zipfile
from pathlib import Path
from openpyxl import load_workbook
def near(a,b):
    if b is None:return a in (None,'')
    return isinstance(a,(int,float)) and not isinstance(a,bool) and math.isclose(a,b,rel_tol=1e-8,abs_tol=1e-5)
def recalc(path,root,name):
    d=root/name;d.mkdir()
    r=subprocess.run(['libreoffice',f'-env:UserInstallation={(root/(name+"-profile")).as_uri()}','--headless','--convert-to','xlsx','--outdir',str(d),str(path)],capture_output=True,timeout=180)
    if r.returncode or not (d/path.name).exists():raise RuntimeError('LibreOffice未完成重算')
    return load_workbook(d/path.name,data_only=True)
def main():
    p=argparse.ArgumentParser();p.add_argument('--artifact-dir',type=Path,required=True);p.add_argument('--result',type=Path,required=True);p.add_argument('--reward-file',type=Path,required=True);a=p.parse_args()
    tests=Path(__file__).parent;inp=tests.parent/'data/input_files';rub=json.loads((tests/'rubric.json').read_text());out=a.artifact_dir/'hr_review.xlsx';scores={c['id']:0. for c in rub['criteria']};ev={}
    def mark(k,x,e):scores[k]=sum(bool(v) for v in x)/len(x) if x else 0.;ev[k]=e
    w=None
    try:
        if out.is_file():w=load_workbook(out)
    except (OSError,ValueError,KeyError,zipfile.BadZipFile):pass
    if w:
        mark('R001',['联合底稿' in w and w['联合底稿'].max_row>1],'有效且非空员工业务底稿')
        tables={}
        for f in inp.glob('*.csv'):
            with f.open(encoding='utf-8-sig') as stream:tables[f.name]=list(csv.DictReader(stream))
        indices={n:{r['EmployeeNumber']:r for r in rows} for n,rows in tables.items()};master=indices['employee_master_list.csv'];rows=[]
        for ident in sorted(master,key=int):
            m=master[ident];c=indices['compensation_details.csv'][ident];p0=indices['performance_review_2024.csv'][ident];s=indices['employee_satisfaction_survey.csv'][ident];t=indices['commute_and_tenure.csv'][ident]
            rows.append([int(ident),m['Department'],m['JobRole'],int(m['JobLevel']),m['Attrition'],int(c['MonthlyIncome']),int(c['MonthlyRate']),c['OverTime'],int(p0['PerformanceRating']),int(p0['PercentSalaryHike']),int(s['JobSatisfaction']),int(s['WorkLifeBalance']),int(t['YearsAtCompany']),int(t['DistanceFromHome']),int(s['EnvironmentSatisfaction']),int(s['RelationshipSatisfaction']),int(s['JobInvolvement']),{1:'Low',2:'Good',3:'Excellent',4:'Outstanding'}[int(p0['PerformanceRating'])],{1:'Low',2:'Medium',3:'High',4:'Very High'}[int(s['JobSatisfaction'])]])
        def checks(book,sheet,expected):
            if sheet not in book:return [False]
            s=book[sheet];result=[s.max_row==len(expected)+1]
            for r,row in enumerate(expected,2):
                for c,value in enumerate(row,1):
                    got=s.cell(r,c).value;result.append(near(got,value) if isinstance(value,(int,float)) or value is None else got==value)
            return result
        field_checks=checks(w,'联合底稿',rows)
        mark('R002',field_checks[1:] if len(field_checks)>1 else [False],'1470编号逐字段核对；薪资与MonthlyRate、两种量表标签分别判断')
        keys=[w['联合底稿'].cell(r,1).value for r in range(2,w['联合底稿'].max_row+1)] if '联合底稿' in w else []
        # 关联行数/主键占本项30%，不能被数万字段稀释成无关紧要的一次检查。
        scores['R002']=.7*scores['R002']+.3*(keys==[r[0] for r in rows])
        keyaudit=[]
        for name,table in sorted(tables.items()):
            keys=set(indices[name]);conflicts=0
            for row in table:
                other=master.get(row['EmployeeNumber'])
                if other:
                    for field in set(other)&set(row)-{'EmployeeNumber'}:conflicts+=row[field]!=other[field]
            keyaudit.append([name,len(table),len(keys),len(set(master)-keys),len(keys-set(master)),conflicts])
        mark('R003',checks(w,'主键审计',keyaudit),'八CSV原行/唯一键/缺失/额外/冲突；分部门缺主表键不等于错误')
        deps=sorted({r[1] for r in rows});roles=sorted({r[2] for r in rows})
        def derive(raw):
            mean=lambda rs,k:statistics.mean(r[k] for r in rs) if rs else None
            rate=lambda rs:sum(r[4]=='Yes' for r in rs)/len(rs) if rs else None
            groups=[]
            for dim,keys,k in [('Department',deps,1),('JobRole',roles,2),('Attrition',['No','Yes'],4)]:
                for key in keys:
                    rs=[r for r in raw if r[k]==key];groups.append([dim,key,len(rs),sum(r[4]=='Yes' for r in rs),rate(rs),mean(rs,5),mean(rs,10),mean(rs,8),sum(r[10]<=2 for r in rs)/len(rs),sum(r[7]=='Yes' for r in rs)/len(rs)])
            strata=[]
            for dep in deps:
                for level in range(1,6):
                    for overtime in ['No','Yes']:
                        rs=[r for r in raw if r[1]==dep and r[3]==level and r[7]==overtime]
                        strata.append([dep,level,overtime,len(rs),sum(r[4]=='Yes' for r in rs),rate(rs),mean(rs,5),mean(rs,10)])
            matrix=[]
            for level in range(1,6):
                rs=[r for r in raw if r[3]==level];matrix.append([level,len(rs),len(rs)/len(raw)]+[rate([r for r in raw if r[3]==level and r[1]==dep]) for dep in deps])
            comparison=[]
            for k,dep in enumerate(deps,3):
                original=rate([r for r in raw if r[1]==dep]);adjusted=sum(r[2]*r[k] for r in matrix);comparison.append([dep,original,adjusted,adjusted-original])
            active=[r for r in raw if r[4]=='No'];left=[r for r in raw if r[4]=='Yes']
            metrics=[len(raw),len(left),len(left)/len(raw),statistics.mean(r[4] for r in groups[:3]),sum(r[2]*r[4] for r in groups[:3])/sum(r[2] for r in groups[:3]),mean(active,5),mean(left,5),mean(active,5)-mean(left,5),sum(r[8]>=3 for r in raw),sum(r[10]<=2 for r in raw)]
            return groups,strata,matrix,comparison,metrics
        groups,strata,matrix,comparison,metrics=derive(rows);required=['联合底稿','主键审计','分组统计','分层统计','职级标准化','审查指标','判读与边界']
        def matrix_checks(book,mat,comparison):
            if '职级标准化' not in book:return [False]
            s=book['职级标准化'];result=[]
            for r,row in enumerate(mat,2):result.extend(near(s.cell(r,c).value,x) for c,x in enumerate(row,1))
            for r,row in enumerate(comparison,2):
                result.append(s.cell(r,8).value==row[0]);result.extend(near(s.cell(r,c).value,x) for c,x in enumerate(row[1:],9))
            return result
        with tempfile.TemporaryDirectory(prefix='hr-judge-') as td:
            temp=Path(td)
            try:v=recalc(out,temp,'base')
            except (RuntimeError,subprocess.TimeoutExpired) as e:
                a.result.parent.mkdir(parents=True,exist_ok=True);a.result.write_text(json.dumps({'status':'infrastructure_failure','error':str(e)},ensure_ascii=False));raise SystemExit(2)
            mark('R004',checks(v,'分组统计',groups),'14组按各自分母：人数、离职、收入/满意/绩效、低满意/加班')
            mark('R005',checks(v,'分层统计',strata),'30个部门职级加班层及零样本的空比例，不混成全员分母')
            mark('R006',matrix_checks(v,matrix,comparison),'全员共同5职级权重与3部门标准化，不用各部门自身权重')
            mark('R007',[near(v['审查指标'].cell(r,2).value,x) if '审查指标' in v else False for r,x in enumerate(metrics,2)],'10总控：简单均值、人数加权、状态差异、量表解释')
            coords=[('分组统计',f'{c}{r}') for r in range(2,16) for c in 'CDEFGHIJ']+[('分层统计',f'{c}{r}') for r in range(2,32) for c in 'DEFGH']+[('职级标准化',f'{c}{r}') for r in range(2,7) for c in 'BCDEF']+[('职级标准化',f'{c}{r}') for r in range(2,5) for c in 'IJK']+[('审查指标',f'B{r}') for r in range(2,12)]
            mark('R008',[s in w and w[s][c].data_type=='f' for s,c in coords],'所有统计与权重为原生公式')
            probe=[]
            if all(s in w for s in required) and w['联合底稿']['A2'].value==1:
                q=load_workbook(out);s=q['联合底稿'];s['E2']='No';s['F2']=6993;s['D2']=3
                changed=[r[:] for r in rows];changed[0][4]='No';changed[0][5]=6993;changed[0][3]=3
                first=[s.cell(2,c).value for c in range(1,20)];last=[s.cell(1471,c).value for c in range(1,20)]
                for c in range(1,20):s.cell(2,c).value=last[c-1];s.cell(1471,c).value=first[c-1]
                f=temp/'changed.xlsx';q.save(f)
                try:pv=recalc(f,temp,'probe')
                except (RuntimeError,subprocess.TimeoutExpired) as e:
                    a.result.parent.mkdir(parents=True,exist_ok=True);a.result.write_text(json.dumps({'status':'infrastructure_failure','error':str(e)},ensure_ascii=False));raise SystemExit(2)
                ng,ns,nmat,ncomp,nmetric=derive(changed)
                for sheet,old,new,offset in [('分组统计',groups,ng,0),('分层统计',strata,ns,0),('职级标准化',matrix,nmat,0),('职级标准化',comparison,ncomp,7)]:
                    for r,(before,after) in enumerate(zip(old,new),2):
                        for c,(b,x) in enumerate(zip(before,after),1):
                            if isinstance(x,(int,float)) and not near(b,x):probe.append(near(pv[sheet].cell(r,c+offset).value,x))
                probe.extend(near(pv['审查指标'].cell(r,2).value,x) for r,(b,x) in enumerate(zip(metrics,nmetric),2) if not near(b,x))
            mark('R009',probe,'一人的离职/收入/职级变更及底稿重排后，跨分组与标准化联动')
        lines=[' '.join(str(x or '') for x in row) for row in w['判读与边界'].iter_rows(min_row=2,max_col=4,values_only=True)] if '判读与边界' in w else []
        wants=[['1470','2940','主表'],['Excellent','High','MonthlyRate','MonthlyIncome'],['权重','全员','构成'],['不能','因果','小样本']]
        mark('R010',[any(all(t in line for t in tokens) and len(line)>35 for line in lines) for tokens in wants]+[any(s._charts for s in w),not w._external_links,all(s in w and w[s].freeze_panes and w[s].auto_filter.ref for s in required)],'关键解释边界与原生图表；不评估个人就业风险')
    score=round(sum(c['weight']*scores[c['id']] for c in rub['criteria']),6)
    if scores['R001']<1:score=0
    result={'task_id':rub['task_id'],'rubric_version':rub['rubric_version'],'judge_version':'hr-attrition-stratification-v1','reward':score,'passed':score>=rub['pass_threshold'],'criterion_results':[{'criterion_id':k,'score':round(v,6),'evidence':ev.get(k,'未满足'),'method':'deterministic'} for k,v in scores.items()]}
    a.result.parent.mkdir(parents=True,exist_ok=True);a.result.write_text(json.dumps(result,ensure_ascii=False,indent=2));a.reward_file.parent.mkdir(parents=True,exist_ok=True);a.reward_file.write_text(str(score)+'\n')
    raise SystemExit(0 if result['passed'] else 1)
if __name__=='__main__':main()
