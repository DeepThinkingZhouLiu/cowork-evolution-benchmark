"""仅针对薪酬核对题的参考解、错误样例与重复判卷检查，不构造其他题。"""
import argparse
import json
import shutil
import subprocess
import sys
from pathlib import Path
from openpyxl import Workbook,load_workbook

def main():
    p=argparse.ArgumentParser();p.add_argument('--output',type=Path,required=True);a=p.parse_args()
    root=Path(__file__).resolve().parents[1];out=a.output.resolve();out.mkdir(parents=True,exist_ok=True)
    if (out/'validation.json').exists():raise FileExistsError('保留旧证据，请指定新验证目录')
    fresh=out/'fresh';subprocess.run([sys.executable,str(root/'solution/solve.py'),'--input-dir',str(root/'data/input_files'),'--output-dir',str(fresh)],check=True)
    cases={'fresh':fresh,'golden':root/'solution/Golden'}
    for label in ('empty','corrupt','blank','unchanged','wrong_salary','missing_report','constant_formulas','no_chart'):
        target=out/label;target.mkdir();cases[label]=target
        if label=='empty':continue
        path=target/'payroll_review.xlsx'
        if label=='corrupt':path.write_bytes(b'not an xlsx');continue
        if label=='blank':Workbook().save(path);continue
        if label=='unchanged':shutil.copy2(next((root/'data/input_files').glob('1-*.xlsx')),path);continue
        w=load_workbook(fresh/'payroll_review.xlsx')
        if label=='wrong_salary':w['年度底稿']['D3']=1
        if label=='missing_report':del w['报告底稿']
        if label=='no_chart':w['部门核对']._charts=[]
        if label=='constant_formulas':
            # 用重算后的正确值伪装公式，检查行为回放能否识别没有依赖的答案。
            cached=out/'cached';cached.mkdir()
            subprocess.run(['libreoffice',f'-env:UserInstallation={(out/"cache-profile").as_uri()}','--headless','--convert-to','xlsx','--outdir',str(cached),str(fresh/'payroll_review.xlsx')],check=True,capture_output=True)
            values=load_workbook(cached/'payroll_review.xlsx',data_only=True)
            for s in w:
                for row in s:
                    for cell in row:
                        if cell.data_type=='f':cell.value='='+str(values[s.title][cell.coordinate].value)
        w.save(path)
    scores={}
    for label,folder in cases.items():
        evidence=out/'judges'/label;evidence.mkdir(parents=True)
        cmd=[sys.executable,str(root/'tests/judge.py'),'--artifact-dir',str(folder),'--result',str(evidence/'result.json'),'--reward-file',str(evidence/'reward.txt')]
        with (evidence/'judge.log').open('w') as log:r=subprocess.run(cmd,stdout=log,stderr=subprocess.STDOUT,timeout=300)
        assert r.returncode in (0,1),(label,r.returncode)
        scores[label]=json.loads((evidence/'result.json').read_text())['reward'];print(label,scores[label],flush=True)
    assert scores['fresh']==scores['golden']==1
    assert all(scores[x]==0 for x in ('empty','corrupt','blank','unchanged'))
    assert 0<scores['wrong_salary']<1 and 0<scores['missing_report']<1
    assert scores['constant_formulas']<.8 and 0<scores['no_chart']<1
    repeat=[]
    for i in range(5):
        evidence=out/'repeat'/str(i);evidence.mkdir(parents=True)
        with (evidence/'judge.log').open('w') as log:subprocess.run([sys.executable,str(root/'tests/judge.py'),'--artifact-dir',str(fresh),'--result',str(evidence/'result.json'),'--reward-file',str(evidence/'reward.txt')],stdout=log,stderr=subprocess.STDOUT,check=True,timeout=300)
        repeat.append(json.loads((evidence/'result.json').read_text())['reward'])
        print('repeat',i,repeat[-1],flush=True)
    assert repeat==[1]*5
    (out/'validation.json').write_text(json.dumps({'task':'workspace-payroll-reconciliation','scores':scores,'repeated_rewards':repeat,'all_assertions_passed':True},ensure_ascii=False,indent=2))
if __name__=='__main__':main()
