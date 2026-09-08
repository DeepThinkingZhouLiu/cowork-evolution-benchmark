"""本题专属解：分开年度/月度版本，重建薪酬证据和情景公式。"""
import argparse
import re
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.workbook.properties import CalcProperties

def main():
    p=argparse.ArgumentParser();p.add_argument('--input-dir',type=Path,required=True);p.add_argument('--output-dir',type=Path,required=True);a=p.parse_args()
    w=Workbook();w.remove(w.active)
    annual=w.create_sheet('年度底稿');annual.append(['版本','年份','人数','年度薪酬','业绩','原列月均人均薪酬','来源定位'])
    src=next(a.input_dir.glob('1-*.xlsx'));sw=load_workbook(src,data_only=True)
    for version,s in zip(('A','B'),sw):
        for n in range(7,12):annual.append([version,s[f'D{n}'].value,s[f'E{n}'].value,s[f'F{n}'].value,s[f'J{n}'].value,s[f'I{n}'].value,f'{src.name}/{s.title}!{n}'])
    month=w.create_sheet('月度底稿');month.append(['版本','部门','人数','工资','来源定位'])
    for pattern,labels,start in [('3-*.xlsx',('图A','图B'),4),('4-*.xlsx',('明细A','明细B'),3)]:
        src=next(a.input_dir.glob(pattern));sw=load_workbook(src,data_only=True)
        for label,s in zip(labels,sw):
            for n in range(start,start+7):month.append([label,s[f'B{n}'].value,s[f'C{n}'].value,s[f'D{n}'].value,f'{src.name}/{s.title}!{n}'])
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with ZipFile(next(a.input_dir.glob('5-*.docx'))) as z:doc=ET.fromstring(z.read('word/document.xml'))
    text='\n'.join(''.join(t.text or '' for t in el.findall('.//w:t',ns)) for el in doc.findall('.//w:p',ns))
    people=int(re.search(r'拥有员工(\d+)人',text)[1]);total=float(re.search(r'月薪酬总额([\d,]+)元',text)[1].replace(',',''));mean=float(re.search(r'人均月薪([\d,]+)元',text)[1].replace(',',''));fund=float(re.search(r'学习基金（人均([\d,]+)元',text)[1].replace(',',''))
    report=w.create_sheet('报告底稿');report.append(['指标','原文数值','单位','来源章节'])
    for label,value,chapter in [('报告人数',people,'公司背景与报告目的'),('报告月总额',total,'2.1.1 整体薪酬分布'),('报告月均薪酬',mean,'2.1.1 整体薪酬分布'),('学习基金每人年额',fund,'短期措施/福利体系升级')]:report.append([label,value,'人' if label=='报告人数' else '元',f'5-薪酬分析报告.docx/{chapter}'])
    tiers=[]
    for table in doc.findall('.//w:tbl',ns):
        rows=[[''.join(t.text or '' for t in c.findall('.//w:t',ns)) for c in r.findall('w:tc',ns)] for r in table.findall('w:tr',ns)]
        if rows and rows[0][:3]==['职级','人数','平均月薪']:tiers=rows[1:]
    scenario=w.create_sheet('职级情景');scenario.append(['职级','人数','原月薪','假设涨幅','原月成本','新月成本','月增量'])
    for n,row in enumerate(tiers,2):
        rate=.10 if row[0]=='中级' else .18 if row[0]=='专家' else 0
        scenario.append([row[0],int(row[1]),float(row[2].replace(',','').replace('元','')),rate,f'=B{n}*C{n}',f'=E{n}*(1+D{n})',f'=F{n}-E{n}'])
    metrics=w.create_sheet('核对指标');metrics.append(['指标','计算结果','来源/口径'])
    entries=[
        ('A版2021至2024薪酬合计','=SUMIFS(年度底稿!D2:D11,年度底稿!A2:A11,"A",年度底稿!B2:B11,">=2021",年度底稿!B2:B11,"<=2024")','A版/2021—2024'),
        ('A版同期原列人均薪酬算术均值','=AVERAGEIFS(年度底稿!F2:F11,年度底稿!A2:A11,"A",年度底稿!B2:B11,">=2021",年度底稿!B2:B11,"<=2024")','原列按年算术平均，非加权平均'),
        ('图A人数','=SUMIF(月度底稿!A2:A29,"图A",月度底稿!C2:C29)','图A部门合计'),
        ('图B人数','=SUMIF(月度底稿!A2:A29,"图B",月度底稿!C2:C29)','图B部门合计'),
        ('图A工资','=SUMIF(月度底稿!A2:A29,"图A",月度底稿!D2:D29)','图A部门合计'),
        ('图B工资','=SUMIF(月度底稿!A2:A29,"图B",月度底稿!D2:D29)','图B部门合计'),
        ('B减A人数差','=B5-B4','版本差异，非人员流动'),
        ('报告月总额除人数','=报告底稿!B3/报告底稿!B2','报告内部算术复核'),
        ('报告公布均薪减重算均薪','=报告底稿!B4-B9','公布减重算'),
        ('职级表月成本合计','=SUM(职级情景!E2:E5)','人数乘月薪后合计'),
        ('职级月成本减报告月总额','=B11-报告底稿!B3','内部口径差异待核实'),
        ('情景年度新增工资','=SUM(职级情景!G2:G5)*12','全年静态情景，非批准支出'),
        ('报告口径年度学习基金','=报告底稿!B2*报告底稿!B5','仅用报告人数')]
    for row in entries:metrics.append(row)
    checks=w.create_sheet('部门核对');checks.append(['部门','A版人数差','A版工资差','B版人数差','B版工资差','B减A人数差','图A工资','图B工资'])
    for n,dept in enumerate([month.cell(i,2).value for i in range(2,9)],2):
        def s(label,col):return f'SUMIFS(月度底稿!{col}2:{col}29,月度底稿!A2:A29,"{label}",月度底稿!B2:B29,A{n})'
        checks.append([dept,'='+s('图A','C')+'-'+s('明细A','C'),'='+s('图A','D')+'-'+s('明细A','D'),'='+s('图B','C')+'-'+s('明细B','C'),'='+s('图B','D')+'-'+s('明细B','D'),'='+s('图B','C')+'-'+s('图A','C'),'='+s('图A','D'),'='+s('图B','D')])
    conclusion=w.create_sheet('审查结论');conclusion.append(['问题','证据定位','数值依据','判断','后续核验动作'])
    for row in [
        ['年度版本','1-公司_年度_人力资源薪酬分析表.xlsx/表格与表格 (2)','A:2020—2024；B:2021—2025','不同版本不可合并，不能认定后缀版获批','请HR确认年度口径和批准版本'],
        ['月度版本','3-公司_月度_薪酬分析图表.xlsx 与 4-公司_月度_薪酬分析.xlsx','A=182人，B=183人；工资相同','每版本图表/明细一致，版本之间人数不同，未说明月份','确认适用月份和人员范围'],
        ['跨期比较','年度底稿/报告概述/分析方法与数据来源','报告131人、2025年1—6月；年度2021—2024','期间和人员范围不同，不可直接比较，不宣称增长率','核实报告样本与年度在职范围'],
        ['均薪算术','报告底稿/2.1.1','1836200/131，与公布14020不同','差额待核实，不能擅自更正原文','核实四舍五入或发放范围'],
        ['职级成本','报告职级薪酬分布与市场对标/2.1.1','2254500与1836200差418300','报告内部不一致，待核实，不混算工资预算','确认职级均薪是否同一基期'],
        ['情景假设','报告短期措施/职级情景','中级10%、专家18%演算','岗位层级到职级的假设映射，非已批准支出','批准前核实对象和实施月份']]:conclusion.append(row)
    chart=BarChart();chart.title='两个版本部门工资核对';chart.width=23;chart.height=10;chart.add_data(Reference(checks,min_col=7,max_col=8,min_row=1,max_row=8),titles_from_data=True);chart.set_categories(Reference(checks,min_col=1,min_row=2,max_row=8));checks.add_chart(chart,'A12')
    w.calculation=CalcProperties(calcMode='auto',fullCalcOnLoad=True,forceFullCalc=True)
    for s in w:
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
        for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='23465B')
        for col in s.columns:
            s.column_dimensions[col[0].column_letter].width=25 if col[0].column<5 else 38
            for c in col[1:]:
                c.alignment=Alignment(vertical='top',wrap_text=True)
                if c.data_type=='f' or isinstance(c.value,(int,float)):c.number_format='#,##0.00;[Red](#,##0.00)'
        for row in range(2,s.max_row+1):s.row_dimensions[row].height=34 if s.title!='审查结论' else 85
    for row in range(2,6):scenario.cell(row,4).number_format='0.0%'
    annual.column_dimensions['G'].width=65;month.column_dimensions['E'].width=65
    report.column_dimensions['D'].width=65
    for s in (annual,month,report):
        for row in range(2,s.max_row+1):s.row_dimensions[row].height=48
    for s in w:
        s.sheet_properties.pageSetUpPr.fitToPage=True
        s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3
        s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0
        s.print_title_rows='1:1';s.print_area=s.dimensions
    checks.print_area='A1:H32'
    a.output_dir.mkdir(parents=True,exist_ok=True);w.save(a.output_dir/'payroll_review.xlsx')
if __name__=='__main__':main()
