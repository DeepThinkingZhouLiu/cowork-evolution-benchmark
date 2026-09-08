"""薪酬题独立阅卷：从冻结输入核对，实际重算并扰动公式链；不读 Golden。"""
import argparse
import json
import math
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from openpyxl import load_workbook

def near(a,b):
    return isinstance(a,(int,float)) and not isinstance(a,bool) and math.isclose(a,b,rel_tol=1e-8,abs_tol=.005)

def recalc(path,root,label):
    folder=root/label;folder.mkdir()
    profile=root/(label+'-profile')
    command=['libreoffice',f'-env:UserInstallation={profile.as_uri()}','--headless','--convert-to','xlsx','--outdir',str(folder),str(path)]
    run=subprocess.run(command,capture_output=True,text=True,timeout=90)
    target=folder/path.name
    if run.returncode or not target.exists():raise RuntimeError('LibreOffice重算未产出文件：'+run.stderr[:300])
    return load_workbook(target,data_only=True)

def main():
    p=argparse.ArgumentParser();p.add_argument('--artifact-dir',type=Path,required=True);p.add_argument('--result',type=Path,required=True);p.add_argument('--reward-file',type=Path,required=True);a=p.parse_args()
    tests=Path(__file__).parent;inp=tests.parent/'data/input_files'
    rubric=json.loads((tests/'rubric.json').read_text());out=a.artifact_dir/'payroll_review.xlsx'
    scores={x['id']:0. for x in rubric['criteria']};ev={k:'未交付或未满足该项' for k in scores}
    def mark(cid,checks,detail):
        scores[cid]=sum(bool(x) for x in checks)/len(checks) if checks else 0.
        ev[cid]=detail
    w=None
    if out.is_file() and zipfile.is_zipfile(out):
        try:w=load_workbook(out,data_only=False)
        except (OSError,ValueError,KeyError,zipfile.BadZipFile):pass
    if w:
        scores['R001']=float(any(s in w and w[s].max_row>1 for s in ('年度底稿','月度底稿','报告底稿')))
        ev['R001']='原生工作簿可打开，且至少一个约定底稿非空；全空表或改名原件不得得分'
        src=load_workbook(next(inp.glob('1-*.xlsx')),data_only=True)
        annual=[]
        for v,s in zip(('A','B'),src):
            for row in range(7,12):annual.append([v,s.cell(row,4).value,s.cell(row,5).value,s.cell(row,6).value,s.cell(row,10).value,s.cell(row,9).value])
        monthly=[]
        for prefix,labels,start in [('3-',('图A','图B'),4),('4-',('明细A','明细B'),3)]:
            sw=load_workbook(next(inp.glob(prefix+'*.xlsx')),data_only=True)
            for label,s in zip(labels,sw):
                for row in range(start,start+7):monthly.append([label]+[s.cell(row,c).value for c in (2,3,4)])
        def raw_check(sheet,expected,source_col,extension):
            if sheet not in w:return [False]
            s=w[sheet];checks=[s.max_row==len(expected)+1]
            for i,row in enumerate(expected,2):
                for j,value in enumerate(row,1):
                    actual=s.cell(i,j).value;checks.append(near(actual,value) if isinstance(value,(int,float)) else actual==value)
                locator=str(s.cell(i,source_col).value or '')
                checks.append(extension in locator and ('!' in locator or '行' in locator))
            return checks
        mark('R002',raw_check('年度底稿',annual,7,'.xlsx'),'核对两个年度版本全部十条记录与原值、来源')
        mark('R003',raw_check('月度底稿',monthly,5,'.xlsx'),'核对四个版本共28部门行，不重复汇总')
        # 以下四个报告标量及四个职级数据已逐条与原文核实；固定来源哈希见 provenance。
        report_values=[131,1836200,14020,2000]
        grade_rows=[['初级',35,9500,0],['中级',50,15000,.1],['高级',30,22000,0],['专家',16,32000,.18]]
        checks=[]
        for n,value in enumerate(report_values,2):
            checks.append('报告底稿' in w and near(w['报告底稿'].cell(n,2).value,value))
        for i,row in enumerate(grade_rows,2):
            for j,value in enumerate(row,1):
                actual=w['职级情景'].cell(i,j).value if '职级情景' in w else None
                checks.append(actual==value if isinstance(value,str) else near(actual,value))
        mark('R004',checks,'报告范围与职级原数独立保留；不从月度表猜报告人数')
        expected=[sum(r[3] for r in annual if r[0]=='A' and 2021<=r[1]<=2024),
                  sum(r[5] for r in annual if r[0]=='A' and 2021<=r[1]<=2024)/4,
                  182,183,2443162,2443162,1,1836200/131,14020-1836200/131,
                  2254500,418300,(50*15000*.1+16*32000*.18)*12,131*2000]
        formula_cells=[('核对指标',f'B{i}') for i in range(2,15)]+[('职级情景',f'{c}{i}') for i in range(2,6) for c in 'EFG']+[('部门核对',f'{c}{i}') for i in range(2,9) for c in 'BCDEF']
        required=['年度底稿','月度底稿','报告底稿','职级情景','核对指标','部门核对','审查结论']
        with tempfile.TemporaryDirectory(prefix='payroll-judge-') as td:
            temp=Path(td)
            try:v=recalc(out,temp,'baseline')
            except (subprocess.TimeoutExpired,RuntimeError) as error:
                a.result.parent.mkdir(parents=True,exist_ok=True)
                a.result.write_text(json.dumps({'status':'infrastructure_failure','error':str(error)},ensure_ascii=False))
                raise SystemExit(2)
            mark('R005',[near(v['核对指标'].cell(i,2).value,e) if '核对指标' in v else False for i,e in enumerate(expected,2)],'实际LibreOffice重算13个指标；接受等价公式，不比较公式字符串')
            dept_expected={r[1]:[0,0,0,0,monthly[7+k][2]-r[2]] for k,r in enumerate(monthly[:7])}
            dept_checks=[]
            if '部门核对' in v:
                s=v['部门核对'];seen=[]
                for row in s.iter_rows(min_row=2,max_col=6,values_only=True):
                    key=row[0];seen.append(key)
                    dept_checks.extend(near(x,y) for x,y in zip(row[1:],dept_expected.get(key,[math.nan]*5)))
                dept_checks.append(sorted(str(x) for x in seen)==sorted(dept_expected))
            mark('R006',dept_checks or [False],'七部门按名称比较，五类差额逐格评分')
            mark('R007',[s in w and w[s][c].data_type=='f' for s,c in formula_cells],'检查所有声明的指标、情景、部门核对公式')
            probe=[]
            if all(s in w for s in required):
                # 只改阅卷临时副本，不动候选；先扰动年度值，再改人数/情景参数并打乱部门顺序。
                pw=load_workbook(out)
                if all(isinstance(pw[s][c].value,(int,float)) for s,c in [('年度底稿','D3'),('月度底稿','C2'),('职级情景','D5')]):
                    pw['年度底稿']['D3']=pw['年度底稿']['D3'].value+12000
                    pw['月度底稿']['C2']=pw['月度底稿']['C2'].value+3
                    pw['职级情景']['D5']=.21
                    ms=pw['月度底稿'];left=[ms.cell(2,c).value for c in range(1,6)];right=[ms.cell(8,c).value for c in range(1,6)]
                    for c in range(1,6):ms.cell(2,c,right[c-1]);ms.cell(8,c,left[c-1])
                    pp=temp/'perturbed.xlsx';pw.save(pp)
                    try:pv=recalc(pp,temp,'probe')
                    except (subprocess.TimeoutExpired,RuntimeError) as error:
                        a.result.write_text(json.dumps({'status':'infrastructure_failure','error':str(error)},ensure_ascii=False));raise SystemExit(2)
                    ps=pv['核对指标'];probe=[near(ps['B2'].value,expected[0]+12000),near(ps['B4'].value,185),near(ps['B8'].value,-2),near(ps['B13'].value,(75000+16*32000*.21)*12)]
                    ds={r[0]:r[1:] for r in pv['部门核对'].iter_rows(min_row=2,max_col=6,values_only=True)}
                    probe += [near(ds.get('综合部',[None]*5)[0],3),near(ds.get('综合部',[None]*5)[4],-4),near(ds.get('售后部',[None]*5)[0],0)]
            mark('R008',probe or [False],'底稿/参数变更与部门行重排后的行为回放')
        conclusion_rows=[]
        if '审查结论' in w:
            conclusion_rows=[' '.join(str(c or '') for c in r) for r in w['审查结论'].iter_rows(min_row=2,max_col=5,values_only=True)]
        # 只给低权重证据覆盖分，不假装用关键词评价全部论证质量。
        groups=[('年度',('版本','合并')),('月度',('182','183')),('跨期',('131','范围')),('均薪',('14020','131')),('职级',('418300','核实')),('情景',('假设','批准'))]
        mark('R009',[any(topic in row and all(t in row for t in tokens) and len(row)>35 for row in conclusion_rows) for topic,tokens in groups],'六类明确数值/口径风险及核验动作覆盖；非主观质量评分')
        chart_ok=any(s._charts for s in w)
        native_chart=False
        with zipfile.ZipFile(out) as z:
            for name in z.namelist():
                if name.startswith('xl/charts/chart') and name.endswith('.xml'):
                    body=z.read(name).decode();native_chart |= ('部门核对' in body or '月度底稿' in body) and '<' in body
        mark('R010',[chart_ok and native_chart,all(s in w and w[s].freeze_panes and w[s].auto_filter.ref for s in required),not w._external_links],'可编辑数据图、各表筛选冻结及无外部链接')
    results=[{'criterion_id':k,'score':round(scores[k],6),'evidence':ev[k],'method':'deterministic'} for k in scores]
    score=sum(x['weight']*scores[x['id']] for x in rubric['criteria'])
    for x in rubric['criteria']:
        if x['type']=='hurdle' and scores[x['id']]<1:score=min(score,x['score_cap_on_failure'])
    score=round(max(0,min(1,score)),6)
    payload={'task_id':rubric['task_id'],'rubric_version':rubric['rubric_version'],'judge_version':'payroll-reconciliation-v1','criterion_results':results,'reward':score,'passed':score>=rubric['pass_threshold']}
    a.result.parent.mkdir(parents=True,exist_ok=True);a.result.write_text(json.dumps(payload,ensure_ascii=False,indent=2))
    a.reward_file.parent.mkdir(parents=True,exist_ok=True);a.reward_file.write_text(f'{score:.6f}\n')
    raise SystemExit(0 if payload['passed'] else 1)
if __name__=='__main__':main()
