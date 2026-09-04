#!/usr/bin/env python3
from __future__ import annotations
import argparse, hashlib, json, os
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
OUT='战略转型决策确认表.xlsx'
FILES=['会议录音元数据.json','D1_上午场_转写稿.txt','D2_下午场_转写稿.txt','DAU行为日志_2025Q4-2026Q1.xlsx','NPS问卷数据_2026-02.xlsx','企业客户CRM_2026Q1.xlsx','竞品监控周报_W01-W12.xlsx']
def style(ws,widths):
 ws.freeze_panes='A2'; ws.sheet_view.showGridLines=False
 for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F4E78'); c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
 for c,w in widths.items(): ws.column_dimensions[c].width=w
 if ws.max_row>1:
  ref=f'A1:{ws.cell(1,ws.max_column).column_letter}{ws.max_row}'
  ws.auto_filter.ref=ref
  try:
   t=Table(displayName='T'+''.join(x for x in ws.title if x.isalnum())[:22],ref=ref); t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True); ws.add_table(t)
  except Exception: pass
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def build(inp:Path,outdir:Path):
 outdir.mkdir(parents=True,exist_ok=True); wb=Workbook(); wb.remove(wb.active)
 ws=wb.create_sheet('Decision_Confirmation'); ws.append(['ID','类别','决策内容','提出人','会话/时间戳','原始文件名'])
 dec=[('DEC-01','战略方向','战略方向确定：放弃泛内容平台铺量，聚焦K12数理化三科垂直AI深耕','林晓峰（CEO）','STRAT-2026-D1-AM 00:23:14','D1_上午场_转写稿.txt'),('DEC-02','产品','AI答疑 Q3 上线承诺，建立垂直 AI 能力路线','张远（CPO）','STRAT-2026-D1-PM 00:07:21（metadata）','会议录音元数据.json'),('DEC-03','预算','重新分配预算以支持 AI 转型','王磊（CFO）','STRAT-2026-D1-PM 00:42:08（metadata）','会议录音元数据.json'),('DEC-04','B端业务','启动 B 端业务立项并以学校采购为重点','周燕（CMO）','STRAT-2026-D2-AM 00:55:12（metadata）','会议录音元数据.json'),('DEC-05','技术架构','推进技术架构切换并追加 150 万元预算，提交董事会审批','李明（CTO）/王磊（CFO）','STRAT-2026-D2-PM 01:08:44','D2_下午场_转写稿.txt'),('DEC-06','组织机制','成立竞品快速响应小组，由张远牵头，李明、周燕参与，2 周内提交方案','林晓峰（CEO）','STRAT-2026-D2-PM 00:52:18','D2_下午场_转写稿.txt')]
 for r in dec: ws.append(r)
 style(ws,{'A':12,'B':14,'C':58,'D':22,'E':34,'F':28})
 ws=wb.create_sheet('Risk_Warnings'); ws.append(['ID','风险内容','提出人/会话','原始文件名','关联决策'])
 risks=[('RISK-01','大模型 API 架构迁移至少需要 5 个月，Q3 上线存在风险。','李明；STRAT-2026-D1-AM 01:05:33','D1_上午场_转写稿.txt','DEC-02,DEC-05'),('RISK-02','B端客户流失主因是缺失AI功能，若不快速上线将加剧流失。','CRM/NPS 交叉证据；STRAT-2026-D2-AM','企业客户CRM_2026Q1.xlsx；NPS问卷数据_2026-02.xlsx','DEC-04'),('RISK-03','竞品功能更完整且快速迭代，B端窗口期可能缩短。','CRM/竞品周报；STRAT-2026-D2-PM','企业客户CRM_2026Q1.xlsx；竞品监控周报_W01-W12.xlsx','DEC-06')]
 for r in risks: ws.append(r)
 style(ws,{'A':12,'B':62,'C':34,'D':48,'E':18})
 ws=wb.create_sheet('OKR_Revisions'); ws.append(['ID','目标','修订内容','提出人/会话','原始文件名'])
 okrs=[('OKR-01-KR1','O1','聚焦 K12 数理化三科垂直 AI，完成垂直路线图','林晓峰；D1 AM','D1_上午场_转写稿.txt'),('OKR-01-KR2','O1','AI 答疑按 Q3 计划上线并建立三科能力基线','张远；D1 PM metadata','会议录音元数据.json'),('OKR-01-KR3','O1','难题退出率从 43% 降低至 25% 以下','张远；D1 AM 00:04:22','D1_上午场_转写稿.txt；DAU行为日志_2025Q4-2026Q1.xlsx'),('OKR-02-KR1','O2','提升答疑解释的理解率与用户继续学习率','周燕；NPS/DAU evidence','NPS问卷数据_2026-02.xlsx；DAU行为日志_2025Q4-2026Q1.xlsx'),('OKR-02-KR2','O2','建立 AI 答案准确性验证和性能改进机制','李明；D1 AM','D1_上午场_转写稿.txt'),('OKR-03-KR1','O3','AI 研发投入翻倍（原目标为保持现有研发投入）','王磊；D1 AM','D1_上午场_转写稿.txt'),('OKR-03-KR2','O3','完成架构迁移成本拆分并控制 Q3 上线风险','李明/王磊；D2 PM','D2_下午场_转写稿.txt'),('OKR-04-KR1','O4','B端增长从当前 23% 向 35% 目标提升，建立学校采购渠道','周燕；D2 AM metadata','企业客户CRM_2026Q1.xlsx；会议录音元数据.json')]
 for r in okrs: ws.append(r)
 style(ws,{'A':18,'B':10,'C':58,'D':32,'E':54})
 ws=wb.create_sheet('Evidence_Metrics'); ws.append(['指标','数值','单位','来源文件','说明'])
 mets=[('难题退出率',0.43,'%','DAU行为日志_2025Q4-2026Q1.xlsx；D1_上午场_转写稿.txt','2026-02 路径漏斗；高于竞品 18 个百分点'),('竞品差距',18,'个百分点','D1_上午场_转写稿.txt','转写稿明确表述'),('NPS 开放反馈：答疑不好用',147,'条','D1_上午场_转写稿.txt','会议引用 NPS 聚类'),('NPS 开放反馈：定价性价比低',156,'条','D1_上午场_转写稿.txt；NPS问卷数据_2026-02.xlsx','会议与聚类汇总'),('B端客户数',214,'家','企业客户CRM_2026Q1.xlsx','CRM 汇总'),('B端增速',0.23,'%','企业客户CRM_2026Q1.xlsx','低于 35% 目标'),('猿辅导 AI 月活',3000000,'人/月','竞品监控周报_W01-W12.xlsx；D2_下午场_转写稿.txt','W12 与会议引用'),('猿辅导 B端 BD 团队',80,'人','竞品监控周报_W01-W12.xlsx；D1_上午场_转写稿.txt','学校白名单采购流程')]
 for r in mets: ws.append(r)
 ws['B2'].number_format='0.0%'; ws['B7'].number_format='0.0%'; style(ws,{'A':28,'B':16,'C':14,'D':54,'E':42})
 ws=wb.create_sheet('90_Day_Action_Plan'); ws.append(['阶段','关键行动','负责人','依据'])
 acts=[('Q2','完成 AI 答疑产品与架构方案、迁移成本拆分','李明（CTO）','D1_上午场_转写稿.txt；D2_下午场_转写稿.txt'),('Q2','成立竞品快速响应小组并提交差异化方案（2周内）','张远（牵头）、李明、周燕','D2_下午场_转写稿.txt'),('Q3','按计划上线 AI 答疑并跟踪难题退出率','张远（CPO）','会议录音元数据.json；DAU行为日志_2025Q4-2026Q1.xlsx'),('Q3','推进学校白名单采购与 B端增长复盘','周燕（CMO）','企业客户CRM_2026Q1.xlsx；竞品监控周报_W01-W12.xlsx')]
 for r in acts: ws.append(r)
 style(ws,{'A':12,'B':64,'C':30,'D':54})
 ws=wb.create_sheet('File_Relationship_Graph'); ws.append(['记录类型','From','To','关系/说明'])
 nodes=FILES+['决策确认表.xlsx','文件关系图谱.json','战略转型决策报告.pptx']
 for n in nodes: ws.append(['NODE',n,'','source or compressed logical output'])
 edges=[('D1_上午场_转写稿.txt','D2_下午场_转写稿.txt','session continuity'),('D1_上午场_转写稿.txt','会议录音元数据.json','session metadata'),('D2_下午场_转写稿.txt','会议录音元数据.json','session metadata'),('D1_上午场_转写稿.txt','决策确认表.xlsx','decision evidence'),('D2_下午场_转写稿.txt','决策确认表.xlsx','decision evidence'),('会议录音元数据.json','决策确认表.xlsx','timestamp/person anchors'),('NPS问卷数据_2026-02.xlsx','决策确认表.xlsx','NPS evidence'),('DAU行为日志_2025Q4-2026Q1.xlsx','决策确认表.xlsx','DAU evidence'),('企业客户CRM_2026Q1.xlsx','决策确认表.xlsx','CRM evidence'),('竞品监控周报_W01-W12.xlsx','决策确认表.xlsx','competitor evidence'),('D1_上午场_转写稿.txt','文件关系图谱.json','lineage'),('D2_下午场_转写稿.txt','文件关系图谱.json','lineage'),('DAU行为日志_2025Q4-2026Q1.xlsx','文件关系图谱.json','lineage'),('NPS问卷数据_2026-02.xlsx','文件关系图谱.json','lineage'),('会议录音元数据.json','文件关系图谱.json','lineage'),('竞品监控周报_W01-W12.xlsx','文件关系图谱.json','lineage'),('企业客户CRM_2026Q1.xlsx','文件关系图谱.json','lineage'),('决策确认表.xlsx','文件关系图谱.json','output lineage'),('D1_上午场_转写稿.txt','战略转型决策报告.pptx','report evidence'),('D2_下午场_转写稿.txt','战略转型决策报告.pptx','report evidence'),('DAU行为日志_2025Q4-2026Q1.xlsx','战略转型决策报告.pptx','report evidence'),('NPS问卷数据_2026-02.xlsx','战略转型决策报告.pptx','report evidence'),('企业客户CRM_2026Q1.xlsx','战略转型决策报告.pptx','report evidence'),('竞品监控周报_W01-W12.xlsx','战略转型决策报告.pptx','report evidence'),('决策确认表.xlsx','战略转型决策报告.pptx','compressed output lineage'),('文件关系图谱.json','战略转型决策报告.pptx','compressed output lineage')]
 for e in edges: ws.append(['EDGE',*e])
 ws.append(['SUMMARY','conclusions','decisions:6; risks:3; okr_revisions:8','compressed one_xlsx adaptation'])
 for r in [('RISK-01','D1_上午场_转写稿.txt'),('RISK-02','企业客户CRM_2026Q1.xlsx；NPS问卷数据_2026-02.xlsx'),('RISK-03','企业客户CRM_2026Q1.xlsx；竞品监控周报_W01-W12.xlsx')]: ws.append(['TRACEABILITY',r[0],r[1],'risk source'])
 for i in range(1,9): ws.append(['TRACEABILITY',f'OKR-{i}','OKR_Revisions row '+str(i+1),'source listed in OKR sheet'])
 style(ws,{'A':18,'B':42,'C':54,'D':32})
 ws=wb.create_sheet('Source_Notes'); ws.append(['项目','记录'])
 ws.append(['Metadata revision','Workspace-Bench-Lite 60b08b1cc2e8054afbc3ca2160d37876b4f0765c；task 386'])
 ws.append(['Workspace archive','Workspace-Bench-Workspaces e245d63bfa20cfdb708cd8e78145ffb087155857；README 声明 Apache-2.0'])
 ws.append(['Adaptation boundary','原始输出含 PPTX/JSON/XLSX；矩阵 one_xlsx，故压缩为一个 XLSX 的决策表、指标表和关系图谱。'])
 ws.append(['Missing source gap','metadata 引用 D1 下午场、D2 上午场和若干音频，但本 task data_manifest 未提供；不臆造其内容。'])
 for fn in FILES:
  p=inp/fn; ws.append([fn,sha(p) if p.exists() else 'MISSING'])
 style(ws,{'A':32,'B':100})
 wb.calculation.fullCalcOnLoad=True; wb.calculation.forceFullCalc=True; wb.calculation.calcMode='auto'; wb.save(outdir/OUT); wb.close()
if __name__=='__main__':
 ap=argparse.ArgumentParser(); ap.add_argument('--input-dir',type=Path,default=Path(os.getenv('INPUT_DIR','/app/data/input_files'))); ap.add_argument('--output-dir',type=Path,default=Path(os.getenv('OUTPUT_DIR','/app/output'))); a=ap.parse_args(); build(a.input_dir,a.output_dir)
