#!/usr/bin/env python3
from __future__ import annotations
import argparse,json,zipfile
from pathlib import Path
from openpyxl import load_workbook
OUT='战略转型决策确认表.xlsx'
def c(i,ok,e): return {'criterion_id':i,'score':1.0 if ok else 0.0,'evidence':e,'method':'deterministic'}
def rows(ws): return [r for r in ws.iter_rows(min_row=2,values_only=True) if any(x not in (None,'') for x in r)]
def main():
 ap=argparse.ArgumentParser(); ap.add_argument('--artifact-dir',type=Path,required=True); ap.add_argument('--result',type=Path,required=True); ap.add_argument('--reward-file',type=Path,required=True); a=ap.parse_args(); out=a.artifact_dir/OUT; rs=[]
 try:
  good=out.is_file() and zipfile.is_zipfile(out); rs.append(c('R001',good,'valid xlsx' if good else 'missing/corrupt'))
  if not good: raise ValueError('invalid output')
  wb=load_workbook(out,data_only=True); required=['Decision_Confirmation','Risk_Warnings','OKR_Revisions','Evidence_Metrics','90_Day_Action_Plan','File_Relationship_Graph','Source_Notes']; exist=all(n in wb.sheetnames for n in required)
  ds=rows(wb['Decision_Confirmation']) if exist else []; dmap={str(r[0]):' '.join(str(x or '') for x in r) for r in ds}; dec_ok=exist and set(dmap)=={f'DEC-0{i}' for i in range(1,7)} and 'K12' in dmap['DEC-01'] and '数理化' in dmap['DEC-01'] and '竞品快速响应小组' in dmap['DEC-06'] and '张远' in dmap['DEC-06']; rs.append(c('R002',dec_ok,f'decisions={len(ds)}'))
  rr=rows(wb['Risk_Warnings']) if exist else []; rmap={str(r[0]):' '.join(str(x or '') for x in r) for r in rr}; risk_ok=set(rmap)=={'RISK-01','RISK-02','RISK-03'} and 'B端客户流失主因是缺失AI功能' in rmap.get('RISK-02',''); rs.append(c('R003',risk_ok,f'risks={len(rr)}'))
  ors=rows(wb['OKR_Revisions']) if exist else []; otext=' '.join(' '.join(str(x or '') for x in r) for r in ors); okr_ok=len(ors)==8 and {'O1','O2','O3','O4'}<=set(str(r[1]) for r in ors) and '43%' in otext and '25%' in otext and 'AI 研发投入翻倍' in otext; rs.append(c('R004',okr_ok,f'okrs={len(ors)}'))
  mr=rows(wb['Evidence_Metrics']) if exist else []; mtext=' '.join(' '.join(str(x or '') for x in r) for r in mr); metric_ok=all(x in mtext for x in ['0.43','18','147','156','214','0.23','3000000','80']); rs.append(c('R005',metric_ok,'key source values'))
  gr=rows(wb['File_Relationship_Graph']) if exist else []; node=sum(1 for r in gr if r[0]=='NODE'); edge=sum(1 for r in gr if r[0]=='EDGE'); trace=sum(1 for r in gr if r[0]=='TRACEABILITY'); graph_ok=node>=10 and edge>=19 and trace>=11; rs.append(c('R006',graph_ok,f'nodes={node}; edges={edge}; trace={trace}'))
  ar=rows(wb['90_Day_Action_Plan']) if exist else []; ntext=' '.join(' '.join(str(x or '') for x in r) for r in rows(wb['Source_Notes'])) if exist else ''; source_ok=any('Q2' in str(r[0]) for r in ar) and any('Q3' in str(r[0]) for r in ar) and '60b08b1' in ntext and 'Apache-2.0' in ntext and 'Missing source gap' in ntext; rs.append(c('R007',source_ok,'action and notes'))
  format_ok=all(wb[n].freeze_panes=='A2' and bool(wb[n].tables) for n in required); rs.append(c('R008',format_ok,'tables and frozen headers')); wb.close()
  weights=[.1,.18,.14,.16,.12,.14,.08,.08]; reward=sum(x['score']*w for x,w in zip(rs,weights)); payload={'task_id':'cowork-evo/workspace-strategy-long-horizon','rubric_version':'workspace-strategy-long-horizon-v0.1','criterion_results':rs,'reward':reward,'passed':all(x['score']==1 for x in rs[:6]) and reward>=.8}
 except Exception as e: payload={'task_id':'cowork-evo/workspace-strategy-long-horizon','criterion_results':rs+[c('R001',False,str(e))],'reward':0,'passed':False}
 a.result.parent.mkdir(parents=True,exist_ok=True); a.result.write_text(json.dumps(payload,ensure_ascii=False,indent=2)); a.reward_file.parent.mkdir(parents=True,exist_ok=True); a.reward_file.write_text(f"{payload['reward']:.6f}\n"); raise SystemExit(0 if payload['passed'] else 1)
if __name__=='__main__': main()
