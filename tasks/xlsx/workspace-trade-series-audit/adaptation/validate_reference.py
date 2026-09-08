"""贸易题逐项反例回放：窗口、方向、原值保存、加权与公式依赖。"""
import argparse,json,shutil,subprocess,sys
from pathlib import Path
from openpyxl import Workbook,load_workbook
def main():
    p=argparse.ArgumentParser();p.add_argument('--output',type=Path,required=True);a=p.parse_args();out=a.output.resolve();out.mkdir(parents=True,exist_ok=False);root=Path(__file__).resolve().parents[1]
    fresh=out/'fresh';subprocess.run([sys.executable,str(root/'solution/solve.py'),'--input-dir',str(root/'data/input_files'),'--output-dir',str(fresh)],check=True)
    cases={'fresh':fresh,'golden':root/'solution/Golden'}
    for label in ['empty','blank','corrupt','unchanged','overwrite_original','swap_direction','average_shares','include_2016','constant_formulas']:
        folder=out/label;folder.mkdir();cases[label]=folder;path=folder/'trade_audit.xlsx'
        if label=='empty':continue
        if label=='blank':Workbook().save(path);continue
        if label=='corrupt':path.write_bytes(b'not workbook');continue
        if label=='unchanged':shutil.copy2(next((root/'data/input_files').glob('2017*.xlsx')),path);continue
        w=load_workbook(fresh/'trade_audit.xlsx')
        if label=='overwrite_original':
            for r in (9,10,11):w['年度底稿'][f'E{r}']=w['年度底稿'][f'C{r}'].value-w['年度底稿'][f'D{r}'].value
        if label=='swap_direction':
            for r in range(2,50):w['商品拆分'][f'B{r}']='进口' if w['商品拆分'][f'B{r}'].value=='出口' else '出口'
        if label=='average_shares':w['窗口指标']['B13']='=AVERAGEIF(\'商品拆分\'!C2:C49,"半导体器件",\'商品拆分\'!D2:D49)'
        if label=='include_2016':w['窗口指标']['B2']="=SUM('年度复核'!F2:F9)+'年度底稿'!B3"
        if label=='constant_formulas':
            cached=out/'cached';cached.mkdir()
            subprocess.run(['libreoffice',f'-env:UserInstallation={(out/"cache-profile").as_uri()}','--headless','--convert-to','xlsx','--outdir',str(cached),str(fresh/'trade_audit.xlsx')],check=True,capture_output=True)
            values=load_workbook(cached/'trade_audit.xlsx',data_only=True)
            for s in w:
                for row in s:
                    for c in row:
                        if c.data_type=='f':c.value='='+str(values[s.title][c.coordinate].value)
        w.save(path)
    scores={}
    for label,folder in cases.items():
        j=out/'judges'/label;j.mkdir(parents=True)
        with (j/'judge.log').open('w') as log:r=subprocess.run([sys.executable,str(root/'tests/judge.py'),'--artifact-dir',str(folder),'--result',str(j/'result.json'),'--reward-file',str(j/'reward.txt')],stdout=log,stderr=subprocess.STDOUT,timeout=300)
        assert r.returncode in (0,1),(label,r.returncode)
        scores[label]=json.loads((j/'result.json').read_text())['reward'];print(label,scores[label],flush=True)
    assert scores['fresh']==scores['golden']==1
    assert all(scores[x]==0 for x in ['empty','blank','corrupt','unchanged'])
    assert all(0<scores[x]<1 for x in ['overwrite_original','swap_direction','average_shares','include_2016'])
    assert scores['constant_formulas']<.8
    repeats=[]
    for i in range(3):
        j=out/f'repeat-{i}';j.mkdir()
        subprocess.run([sys.executable,str(root/'tests/judge.py'),'--artifact-dir',str(fresh),'--result',str(j/'result.json'),'--reward-file',str(j/'reward.txt')],check=True,capture_output=True,timeout=300);repeats.append(json.loads((j/'result.json').read_text())['reward'])
    assert repeats==[1]*3
    (out/'validation.json').write_text(json.dumps({'task':'workspace-trade-series-audit','scores':scores,'repeated_rewards':repeats,'all_assertions_passed':True},ensure_ascii=False,indent=2))
if __name__=='__main__':main()

