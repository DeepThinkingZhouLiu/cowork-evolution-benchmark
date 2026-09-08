"""贸易题参考解：独立保留原值，重算恒等式，按方向拆分商品。"""
import argparse,re
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.chart import LineChart,Reference
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter

def main():
    p=argparse.ArgumentParser();p.add_argument('--input-dir',type=Path,required=True);p.add_argument('--output-dir',type=Path,required=True);a=p.parse_args()
    w=Workbook();w.remove(w.active)
    raw=w.create_sheet('年度底稿');raw.append(['年份','原总额亿元','出口亿元','进口亿元','原差额亿元','出口商品','进口商品','来源'])
    cumulative=w.create_sheet('累计底稿');cumulative.append(['起年','止年','原总额','原出口','原进口','原差额','来源'])
    for f in sorted(a.input_dir.glob('*.xlsx')):
        src=load_workbook(f,data_only=True).active
        for r in (2,3):raw.append([src.cell(r,c).value for c in range(1,8)]+[f'{f.name}!{src.title}!A{r}:G{r}'])
        cumulative.append([src['A2'].value,src['A3'].value]+[src.cell(4,c).value for c in range(2,6)]+[f'{f.name}!{src.title}!A4:E4'])
    commodities=w.create_sheet('商品拆分');commodities.append(['年份','方向','商品','份额','金额亿元','来源'])
    for row in raw.iter_rows(min_row=4,max_row=11,values_only=True):
        year=row[0]
        for direction,text,col in [('出口',row[5],'C'),('进口',row[6],'D')]:
            for name,pct in re.findall(r'([^、]+)\(([\d.]+)%\)',text):
                r=commodities.max_row+1
                commodities.append([year,direction,name,float(pct)/100,f'=SUMIF(\'年度底稿\'!A$2:A$11,A{r},\'年度底稿\'!{col}$2:{col}$11)*D{r}',row[7]+';原'+('F' if direction=='出口' else 'G')+'列'])
    audit=w.create_sheet('年度复核');audit.append(['年份','原总额','出口','进口','原差额','重算总额','重算差额','总额残差','差额残差','总额同比','出口同比','进口同比','进出口比','出口核心金额','进口核心金额'])
    for year in range(2017,2025):
        r=audit.max_row+1
        fields=[year]+[f'=SUMIF(\'年度底稿\'!A$2:A$11,A{r},\'年度底稿\'!{c}$2:{c}$11)' for c in 'BCDE']
        fields += [f'=C{r}+D{r}',f'=C{r}-D{r}',f'=B{r}-F{r}',f'=E{r}-G{r}',
                   f'=F{r}/(SUMIF(\'年度底稿\'!A$2:A$11,A{r}-1,\'年度底稿\'!C$2:C$11)+SUMIF(\'年度底稿\'!A$2:A$11,A{r}-1,\'年度底稿\'!D$2:D$11))-1',
                   f'=C{r}/SUMIF(\'年度底稿\'!A$2:A$11,A{r}-1,\'年度底稿\'!C$2:C$11)-1',
                   f'=D{r}/SUMIF(\'年度底稿\'!A$2:A$11,A{r}-1,\'年度底稿\'!D$2:D$11)-1',f'=C{r}/D{r}',
                   f'=SUMIFS(\'商品拆分\'!E$2:E$49,\'商品拆分\'!A$2:A$49,A{r},\'商品拆分\'!B$2:B$49,"出口")',
                   f'=SUMIFS(\'商品拆分\'!E$2:E$49,\'商品拆分\'!A$2:A$49,A{r},\'商品拆分\'!B$2:B$49,"进口")']
        audit.append(fields)
    ca=w.create_sheet('累计复核');ca.append(['起年','止年','合成总额','合成出口','合成进口','合成差额','差额残差','总额残差'])
    for r in range(2,7):
        ca.append([cumulative.cell(r,1).value,cumulative.cell(r,2).value,None,
                   f'=SUMIFS(\'年度底稿\'!C$2:C$11,\'年度底稿\'!A$2:A$11,">="&A{r},\'年度底稿\'!A$2:A$11,"<="&B{r})',
                   f'=SUMIFS(\'年度底稿\'!D$2:D$11,\'年度底稿\'!A$2:A$11,">="&A{r},\'年度底稿\'!A$2:A$11,"<="&B{r})',
                   f'=D{r}-E{r}',f'=\'累计底稿\'!F{r}-F{r}',f'=\'累计底稿\'!C{r}-C{r}'])
        ca.cell(r,3,f'=D{r}+E{r}')
    metrics=w.create_sheet('窗口指标');metrics.append(['指标','结果'])
    pairs=[('重算总额累计',"SUM('年度复核'!F2:F9)"),('出口累计',"SUM('年度复核'!C2:C9)"),('进口累计',"SUM('年度复核'!D2:D9)"),('重算差额累计',"SUM('年度复核'!G2:G9)"),('原差额累计',"SUM('年度复核'!E2:E9)"),('差额残差累计','B6-B5'),
           ('差额异常年数','COUNTIF(\'年度复核\'!I2:I9,">0.005")+COUNTIF(\'年度复核\'!I2:I9,"<-0.005")'),
           ('2017至2024总额CAGR',"('年度复核'!F9/'年度复核'!F2)^(1/7)-1"),
           ('2024出口机电占比','SUMIFS(\'商品拆分\'!D2:D49,\'商品拆分\'!A2:A49,2024,\'商品拆分\'!B2:B49,"出口",\'商品拆分\'!C2:C49,"机电产品")'),
           ('2024进口半导体器件占比','SUMIFS(\'商品拆分\'!D2:D49,\'商品拆分\'!A2:A49,2024,\'商品拆分\'!B2:B49,"进口",\'商品拆分\'!C2:C49,"半导体器件")'),
           ('进口半导体器件金额累计','SUMIFS(\'商品拆分\'!E2:E49,\'商品拆分\'!B2:B49,"进口",\'商品拆分\'!C2:C49,"半导体器件")'),
           ('跨年加权进口份额','B12/B4')]
    for label,f in pairs:metrics.append([label,'='+f])
    notes=w.create_sheet('审计说明');notes.append(['疑点','定位','数值影响','待核实'])
    for row in [
        ['年度差额异常','2021-2022源表E3、2023-2024源表E2:E3','2022/2023/2024原差额各比出口减进口少2000亿元','保留原值与重算值，请核实出口/进口和差额口径，不能自行宣布底层数正确'],
        ['累计行受影响','两份相关源文件第4行','2021-2022累计差额少2000；2023-2024少4000；窗口合计少6000亿元','累计行不当年度；追踪异常年度的累加影响'],
        ['窗口外基准','2015-2016源文件2016年度','2016只用于2017同比；最终累计仅2017-2024八年','禁止为了同比把2016加入最终窗口；CAGR用七个间隔'],
        ['商品方向','2019-2020文件F3/G3','2020半导体器件20.5%在进口G列，不是出口；设备与器件不同类','逐年逐方向核实；加权占比用金额除进口总额，不平均百分比']]:notes.append(row)
    chart=LineChart();chart.title='出口与进口年度序列（亿元）';chart.add_data(Reference(audit,min_col=3,max_col=4,min_row=1,max_row=9),titles_from_data=True);chart.set_categories(Reference(audit,min_col=1,min_row=2,max_row=9));chart.width=24;chart.height=12;audit.add_chart(chart,'A12')
    for s in w:
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
        for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='24476B')
        for row in s.iter_rows(min_row=2):
            for c in row:
                c.alignment=Alignment(wrap_text=True,vertical='top')
                if c.data_type=='f' or isinstance(c.value,(float,int)):c.number_format='#,##0.00'
        for c in range(1,s.max_column+1):s.column_dimensions[get_column_letter(c)].width=20
        if s.title=='年度底稿':
            for c in 'FGH':s.column_dimensions[c].width=50
            for r in range(2,12):s.row_dimensions[r].height=52
        if s.title=='商品拆分':
            s.column_dimensions['F'].width=70
            for r in range(2,50):s.row_dimensions[r].height=32;s.cell(r,4).number_format='0.00%'
        if s.title=='审计说明':
            for c in 'BCD':s.column_dimensions[c].width=55
            for r in range(2,6):s.row_dimensions[r].height=70
        s.sheet_properties.pageSetUpPr.fitToPage=True;s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0;s.print_title_rows='1:1'
    for r in range(2,10):
        audit.cell(r,1).number_format='0'
        for c in range(10,14):audit.cell(r,c).number_format='0.00%'
    for r in (9,10,11,13):metrics.cell(r,2).number_format='0.00%'
    a.output_dir.mkdir(parents=True,exist_ok=True);w.save(a.output_dir/'trade_audit.xlsx')
if __name__=='__main__':main()
