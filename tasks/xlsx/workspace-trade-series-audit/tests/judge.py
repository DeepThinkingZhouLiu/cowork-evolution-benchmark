"""贸易序列专属Judge：源表独立重算、方向拆分、累计追踪与扰动测试。"""
import argparse,json,math,re,subprocess,tempfile,zipfile
from pathlib import Path
from openpyxl import load_workbook
def near(a,b):return isinstance(a,(int,float)) and not isinstance(a,bool) and math.isclose(a,b,rel_tol=1e-8,abs_tol=.00001)
def calc(path,root,label):
    d=root/label;d.mkdir()
    x=subprocess.run(['libreoffice',f'-env:UserInstallation={(root/(label+"-profile")).as_uri()}','--headless','--convert-to','xlsx','--outdir',str(d),str(path)],capture_output=True,timeout=100)
    if x.returncode or not (d/path.name).exists():raise RuntimeError('重算基础设施未返回文件')
    return load_workbook(d/path.name,data_only=True)
def main():
    p=argparse.ArgumentParser();p.add_argument('--artifact-dir',type=Path,required=True);p.add_argument('--result',type=Path,required=True);p.add_argument('--reward-file',type=Path,required=True);a=p.parse_args()
    tests=Path(__file__).parent;inp=tests.parent/'data/input_files';rub=json.loads((tests/'rubric.json').read_text());out=a.artifact_dir/'trade_audit.xlsx';scores={c['id']:0. for c in rub['criteria']};ev={}
    def mark(k,x,e):scores[k]=sum(bool(v) for v in x)/len(x) if x else 0.;ev[k]=e
    w=None
    try:
        if out.is_file():w=load_workbook(out)
    except (OSError,ValueError,KeyError,zipfile.BadZipFile):pass
    if w:
        mark('R001',['年度底稿' in w and w['年度底稿'].max_row>1],'原生可解析工作簿与业务底稿')
        years={};periods=[]
        for f in sorted(inp.glob('*.xlsx')):
            s=load_workbook(f,data_only=True).active
            for r in (2,3):years[int(s.cell(r,1).value)]=[s.cell(r,c).value for c in range(2,8)]
            periods.append([s['A2'].value,s['A3'].value]+[s.cell(4,c).value for c in range(2,6)])
        def raw(sheet,rows,source_col):
            if sheet not in w:return [False]
            checks=[w[sheet].max_row==len(rows)+1]
            for i,row in enumerate(rows,2):
                for c,value in enumerate(row,1):
                    got=w[sheet].cell(i,c).value;checks.append(near(got,value) if isinstance(value,(int,float)) else got==value)
                checks.append('.xlsx!' in str(w[sheet].cell(i,source_col).value))
            return checks
        mark('R002',raw('年度底稿',[[k]+v for k,v in years.items()],8),'完整保留10年原金额、差额和进出口原文本')
        mark('R003',raw('累计底稿',periods,7),'五个累计期间单列，不混进年度')
        products=[]
        for year in range(2017,2025):
            for direction,idx in [('出口',4),('进口',5)]:
                for token in years[year][idx].split('、'):
                    name,percent=token.rsplit('(',1);products.append([year,direction,name,float(percent.rstrip('%)'))/100])
        mark('R004',raw('商品拆分',products,6),'48商品年-方向-名称-份额精确对应，不混设备/器件')
        def annual_values(y,data,items):
            total,exp,imp,old,*_=data[y];prev=data[y-1]
            core=[sum((exp if direction=='出口' else imp)*share for year,direction,name,share in items if year==y and direction==d) for d in ['出口','进口']]
            return [total,exp,imp,old,exp+imp,exp-imp,total-exp-imp,old-exp+imp,(exp+imp)/(prev[1]+prev[2])-1,exp/prev[1]-1,imp/prev[2]-1,exp/imp]+core
        av={y:annual_values(y,years,products) for y in range(2017,2025)}
        cumulative=[]
        for start,end,total,ex,im,diff in periods:
            exp=sum(years[y][1] for y in range(start,end+1));imp=sum(years[y][2] for y in range(start,end+1));cumulative.append([start,end,exp+imp,exp,imp,exp-imp,diff-exp+imp,total-exp-imp])
        def metrics(data,items):
            ex=sum(data[y][1] for y in range(2017,2025));im=sum(data[y][2] for y in range(2017,2025));old=sum(data[y][3] for y in range(2017,2025))
            semi=sum(data[y][2]*share for y,d,n,share in items if d=='进口' and n=='半导体器件')
            share=lambda direction,name:next(r[3] for r in items if r[:3]==[2024,direction,name])
            return [ex+im,ex,im,ex-im,old,old-ex+im,sum(abs(data[y][3]-data[y][1]+data[y][2])>.005 for y in range(2017,2025)),((data[2024][1]+data[2024][2])/(data[2017][1]+data[2017][2]))**(1/7)-1,share('出口','机电产品'),share('进口','半导体器件'),semi,semi/im]
        mv=metrics(years,products);required=['年度底稿','累计底稿','年度复核','累计复核','商品拆分','窗口指标','审计说明']
        with tempfile.TemporaryDirectory(prefix='trade-judge-') as td:
            temp=Path(td)
            try:v=calc(out,temp,'base')
            except (RuntimeError,subprocess.TimeoutExpired) as error:
                a.result.parent.mkdir(parents=True,exist_ok=True);a.result.write_text(json.dumps({'status':'infrastructure_failure','error':str(error)},ensure_ascii=False));raise SystemExit(2)
            checks=[]
            if '年度复核' in v:
                checks.append(v['年度复核'].max_row==9)
                for r,y in enumerate(range(2017,2025),2):
                    checks.append(v['年度复核'].cell(r,1).value==y)
                    checks.extend(near(v['年度复核'].cell(r,c).value,x) for c,x in enumerate(av[y],2))
            mark('R005',checks,'八年原/重算恒等式、2016基期同比和方向金额')
            checks=[]
            if '累计复核' in v:
                checks.append(v['累计复核'].max_row==6)
                for r,row in enumerate(cumulative,2):checks.extend(near(v['累计复核'].cell(r,c).value,x) for c,x in enumerate(row,1))
            mark('R006',checks,'年度窗口聚合追踪累计差额，不信任原累计差额')
            checks=[near(v['窗口指标'].cell(r,2).value,x) if '窗口指标' in v else False for r,x in enumerate(mv,2)]
            if '商品拆分' in v:
                for r,(y,d,n,share) in enumerate(products,2):checks.append(near(v['商品拆分'].cell(r,5).value,years[y][1 if d=='出口' else 2]*share))
            mark('R007',checks,'12窗口指标与48商品金额；CAGR七间隔、金额加权')
            coords=[('年度复核',f'{c}{r}') for r in range(2,10) for c in 'BCDEFGHIJKLMNO']+[('累计复核',f'{c}{r}') for r in range(2,7) for c in 'CDEFGH']+[('商品拆分',f'E{r}') for r in range(2,50)]+[('窗口指标',f'B{r}') for r in range(2,14)]
            mark('R008',[s in w and w[s][c].data_type=='f' for s,c in coords],'声明所有计算原生公式')
            checks=[]
            if all(s in w for s in required) and near(w['年度底稿']['C11'].value,12285.7):
                q=load_workbook(out);q['年度底稿']['C11']=12385.7;q['商品拆分']['D44']=.568
                s=q['年度底稿'];left=[s.cell(2,c).value for c in range(1,9)];right=[s.cell(11,c).value for c in range(1,9)]
                for c in range(1,9):s.cell(2,c,right[c-1]);s.cell(11,c,left[c-1])
                f=temp/'changed.xlsx';q.save(f)
                try:pv=calc(f,temp,'probe')
                except (RuntimeError,subprocess.TimeoutExpired) as error:
                    a.result.parent.mkdir(parents=True,exist_ok=True);a.result.write_text(json.dumps({'status':'infrastructure_failure','error':str(error)},ensure_ascii=False));raise SystemExit(2)
                newdata={y:row[:] for y,row in years.items()};newdata[2024][1]+=100
                newitems=[row[:] for row in products];newitems[42][3]=.568
                wanted=annual_values(2024,newdata,newitems)
                checks=[near(pv['年度复核'].cell(9,c).value,wanted[c-2]) for c in (3,6,7,8,9,10,11,13,14)]
                newmv=metrics(newdata,newitems)
                checks.extend(near(pv['窗口指标'].cell(i+2,2).value,x) for i,x in enumerate(newmv) if not near(x,mv[i]))
                checks.extend([near(pv['累计复核']['G6'].value,cumulative[-1][6]-100),near(pv['商品拆分']['E44'].value,12385.7*.568)])
            mark('R009',checks,'出口原值与份额改变、年度重排后所有应变结果更新')
        notes=[' '.join(str(x or '') for x in row) for row in w['审计说明'].iter_rows(min_row=2,max_col=4,values_only=True)] if '审计说明' in w else []
        requirements=[['2022','2023','2024','2000'],['累计','4000','6000'],['2016','2017','同比'],['进口','出口','20.5%']]
        mark('R010',[any(all(word in line for word in words) and len(line)>40 for line in notes) for words in requirements]+[any(s._charts for s in w),not w._external_links,all(s in w and w[s].freeze_panes and w[s].auto_filter.ref for s in required)],'四类证据要点和可用性；非主观美学评分')
    reward=round(sum(c['weight']*scores[c['id']] for c in rub['criteria']),6)
    if scores['R001']<1:reward=0
    result={'task_id':rub['task_id'],'rubric_version':rub['rubric_version'],'judge_version':'trade-series-audit-v1','reward':reward,'passed':reward>=rub['pass_threshold'],'criterion_results':[{'criterion_id':k,'score':round(v,6),'evidence':ev.get(k,'未满足'),'method':'deterministic'} for k,v in scores.items()]}
    a.result.parent.mkdir(parents=True,exist_ok=True);a.result.write_text(json.dumps(result,ensure_ascii=False,indent=2));a.reward_file.parent.mkdir(parents=True,exist_ok=True);a.reward_file.write_text(str(reward)+'\n')
    raise SystemExit(0 if result['passed'] else 1)
if __name__=='__main__':main()
