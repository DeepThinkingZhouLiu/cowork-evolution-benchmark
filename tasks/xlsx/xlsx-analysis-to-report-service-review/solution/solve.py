
from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "17365D"
BLUE = "4472C4"
GREEN = "70AD47"
ORANGE = "ED7D31"
RED = "C00000"
WHITE = "FFFFFF"
LIGHT_BLUE = "D9EAF7"
THIN = Side(style="thin", color="B7B7B7")


def header(ws, row=1):
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.font = Font(name="微软雅黑", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ws.row_dimensions[row].height = 28


def body(ws, start_row=2):
    for row in ws.iter_rows(min_row=start_row):
        for cell in row:
            cell.font = Font(name="微软雅黑", size=10, color="222222")
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def finish(ws, header_row=1, freeze=None, autofilter=True):
    header(ws, header_row)
    body(ws, header_row + 1)
    ws.freeze_panes = freeze or f"A{header_row + 1}"
    if autofilter:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.sheet_view.showGridLines = False
    for col in range(1, ws.max_column + 1):
        vals = [str(ws.cell(r, col).value or "") for r in range(1, min(ws.max_row, 80) + 1)]
        width = min(max(max((len(v) for v in vals), default=8) * 1.2 + 2, 10), 30)
        ws.column_dimensions[get_column_letter(col)].width = width


def table(ws, name, ref=None):
    ref = ref or f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    obj = Table(displayName=name, ref=ref)
    obj.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(obj)


def money(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = '¥#,##0.00;[Red]-¥#,##0.00'


def percent(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = "0.00%"


def save(wb, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(path)


def solve(input_dir, output_dir):
    source = load_workbook(input_dir / "客户服务月度数据.xlsx", data_only=True).active
    rows = list(source.iter_rows(min_row=2, values_only=True))
    wb = Workbook()
    data = wb.active
    data.title = "月度数据"
    data.append(["月份", "工单量", "SLA达成率", "CSAT", "流失客户数", "收入", "服务成本", "利润", "利润率"])
    for row in rows:
        profit = row[5] - row[6]
        data.append(list(row) + [profit, profit / row[5]])
    finish(data)
    percent(data, ["C", "I"])
    money(data, ["F", "G", "H"])
    table(data, "ServiceMonthly")

    total_tickets = sum(row[1] for row in rows)
    avg_sla = sum(row[2] for row in rows) / len(rows)
    weighted_csat = sum(row[3] * row[1] for row in rows) / total_tickets
    churn = sum(row[4] for row in rows)
    revenue = sum(row[5] for row in rows)
    cost = sum(row[6] for row in rows)
    profit = revenue - cost
    kpi = wb.create_sheet("KPI摘要")
    kpi.append(["指标", "数值"])
    for row in [
        ["全年工单量", total_tickets], ["平均 SLA 达成率", avg_sla], ["加权平均 CSAT", weighted_csat],
        ["流失客户数", churn], ["全年收入", revenue], ["全年服务成本", cost], ["全年利润", profit], ["全年利润率", profit / revenue],
    ]:
        kpi.append(row)
    finish(kpi)
    for cell in ["B3", "B9"]:
        kpi[cell].number_format = "0.00%"
    kpi["B4"].number_format = "0.00"
    money(kpi, ["B"], 6, 8)

    trend = wb.create_sheet("趋势分析")
    trend.append(["月份", "SLA达成率", "CSAT", "利润", "利润率"])
    for row in rows:
        monthly_profit = row[5] - row[6]
        trend.append([row[0], row[2], row[3], monthly_profit, monthly_profit / row[5]])
    finish(trend)
    percent(trend, ["B", "E"])
    money(trend, ["D"])
    quality_chart = LineChart()
    quality_chart.title = "SLA 与 CSAT 月度趋势"
    quality_chart.add_data(Reference(trend, min_col=2, max_col=3, min_row=1, max_row=trend.max_row), titles_from_data=True)
    quality_chart.set_categories(Reference(trend, min_col=1, min_row=2, max_row=trend.max_row))
    quality_chart.height, quality_chart.width = 8, 15
    trend.add_chart(quality_chart, "G2")
    profit_chart = BarChart()
    profit_chart.title = "月度利润趋势"
    profit_chart.add_data(Reference(trend, min_col=4, min_row=1, max_row=trend.max_row), titles_from_data=True)
    profit_chart.set_categories(Reference(trend, min_col=1, min_row=2, max_row=trend.max_row))
    profit_chart.height, profit_chart.width = 8, 15
    trend.add_chart(profit_chart, "G18")
    save(wb, output_dir / "客户服务经营分析.xlsx")

    lowest_sla = min(rows, key=lambda row: row[2])
    lowest_profit = min(rows, key=lambda row: row[5] - row[6])
    document = Document()
    title = document.add_heading("2026 年客户服务经营分析报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("数据口径：2026 年1–12月合成月度经营数据；本报告不包含外部市场推断。")
    document.add_heading("执行摘要", level=1)
    document.add_paragraph(
        f"全年处理 {total_tickets:,} 个工单，平均 SLA 达成率 {avg_sla:.2%}，"
        f"加权平均 CSAT {weighted_csat:.2f}。全年收入 ¥{revenue:,.2f}，服务成本 ¥{cost:,.2f}，"
        f"利润 ¥{profit:,.2f}，利润率 {profit/revenue:.2%}。"
    )
    table_doc = document.add_table(rows=1, cols=2)
    table_doc.style = "Table Grid"
    table_doc.rows[0].cells[0].text = "指标"
    table_doc.rows[0].cells[1].text = "数值"
    for label, value in [
        ("全年工单量", f"{total_tickets:,}"), ("平均 SLA", f"{avg_sla:.2%}"),
        ("加权 CSAT", f"{weighted_csat:.2f}"), ("流失客户数", f"{churn}"), ("全年利润", f"¥{profit:,.2f}"),
    ]:
        cells = table_doc.add_row().cells
        cells[0].text, cells[1].text = label, value
    document.add_heading("关键发现", level=1)
    document.add_paragraph(f"SLA 最低月份为 {lowest_sla[0]}，达成率 {lowest_sla[2]:.2%}；与全年平均值相比明显偏低。")
    document.add_paragraph(f"利润最低月份为 {lowest_profit[0]}，当月利润 ¥{lowest_profit[5]-lowest_profit[6]:,.2f}。")
    document.add_paragraph(f"全年流失客户 {churn} 个，SLA 较低的 5–7 月同时出现 CSAT 下探，应优先复盘人力排班和工单分流。")
    document.add_heading("改进建议", level=1)
    for text in [
        "1. 针对 5–7 月工单高峰，按周追踪未结工单和超时原因，设置 SLA 预警阈值。",
        "2. 将高频问题建立标准回复和分流规则，以缩短首次响应和解决时间。",
        "3. 对当月低分 CSAT 样本开展分层回访，每月向责任团队反馈前三类根因。",
        "4. 将月度利润与 SLA/CSAT 同时纳入经营会，防止单纯压降成本损害服务质量。",
    ]:
        document.add_paragraph(text)
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.font.name = "微软雅黑"
            run.font.size = Pt(10.5)
    document.save(output_dir / "客户服务经营分析报告.docx")



def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--output-dir", required=True)
    args = parser.parse_args()
    input_dir = Path(args.input_dir).resolve()
    output_dir = Path(args.output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    solve(input_dir, output_dir)


if __name__ == "__main__":
    main()
