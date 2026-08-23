
from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "17365D"
BLUE = "4472C4"
GREEN = "70AD47"
ORANGE = "ED7D31"
RED = "C00000"
WHITE = "FFFFFF"
LIGHT_BLUE = "D9EAF7"
THIN = Side(style="thin", color="B7B7B7")


def header(ws, row=1):
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.font = Font(name="微软雅黑", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ws.row_dimensions[row].height = 28


def body(ws, start_row=2):
    for row in ws.iter_rows(min_row=start_row):
        for cell in row:
            cell.font = Font(name="微软雅黑", size=10, color="222222")
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def finish(ws, header_row=1, freeze=None, autofilter=True):
    header(ws, header_row)
    body(ws, header_row + 1)
    ws.freeze_panes = freeze or f"A{header_row + 1}"
    if autofilter:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.sheet_view.showGridLines = False
    for col in range(1, ws.max_column + 1):
        vals = [str(ws.cell(r, col).value or "") for r in range(1, min(ws.max_row, 80) + 1)]
        width = min(max(max((len(v) for v in vals), default=8) * 1.2 + 2, 10), 30)
        ws.column_dimensions[get_column_letter(col)].width = width


def table(ws, name, ref=None):
    ref = ref or f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    obj = Table(displayName=name, ref=ref)
    obj.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(obj)


def money(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = '¥#,##0.00;[Red]-¥#,##0.00'


def percent(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = "0.00%"


def save(wb, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(path)


def solve(input_dir, output_dir):
    source = load_workbook(input_dir / "增长漏斗周报.xlsx", data_only=True).active
    wb = Workbook()
    data = wb.active
    data.title = "周度数据"
    data.append(["周次", "周开始日", "线索数", "商机数", "订单数", "线索→商机转化率", "商机→订单转化率"])
    for row in source.iter_rows(min_row=2, values_only=True):
        data.append(list(row) + [row[3] / row[2], row[4] / row[3]])
    finish(data)
    percent(data, ["F", "G"])
    table(data, "GrowthFunnel")

    board = wb.create_sheet("图表看板")
    board.append(["指标", "数值"])
    total_leads = sum(data[f"C{r}"].value for r in range(2, data.max_row + 1))
    total_opps = sum(data[f"D{r}"].value for r in range(2, data.max_row + 1))
    total_orders = sum(data[f"E{r}"].value for r in range(2, data.max_row + 1))
    for row in [["累计线索", total_leads], ["累计商机", total_opps], ["累计订单", total_orders], ["整体线索转订单率", total_orders / total_leads]]:
        board.append(row)
    finish(board)
    board["B5"].number_format = "0.00%"

    line = LineChart()
    line.title = "周度漏斗规模趋势"
    line.y_axis.title = "数量"
    line.x_axis.title = "周次"
    line.add_data(Reference(data, min_col=3, max_col=5, min_row=1, max_row=data.max_row), titles_from_data=True)
    line.set_categories(Reference(data, min_col=1, min_row=2, max_row=data.max_row))
    line.height, line.width = 8, 15
    board.add_chart(line, "D2")
    bars = BarChart()
    bars.type = "col"
    bars.title = "周度转化率"
    bars.y_axis.title = "转化率"
    bars.y_axis.numFmt = "0%"
    bars.x_axis.title = "周次"
    bars.add_data(Reference(data, min_col=6, max_col=7, min_row=1, max_row=data.max_row), titles_from_data=True)
    bars.set_categories(Reference(data, min_col=1, min_row=2, max_row=data.max_row))
    bars.height, bars.width = 8, 15
    board.add_chart(bars, "D18")
    save(wb, output_dir / "增长漏斗图表.xlsx")



def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--output-dir", required=True)
    args = parser.parse_args()
    input_dir = Path(args.input_dir).resolve()
    output_dir = Path(args.output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    solve(input_dir, output_dir)


if __name__ == "__main__":
    main()
