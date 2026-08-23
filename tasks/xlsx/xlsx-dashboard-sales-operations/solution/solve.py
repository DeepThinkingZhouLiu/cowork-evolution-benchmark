
from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "17365D"
BLUE = "4472C4"
GREEN = "70AD47"
ORANGE = "ED7D31"
RED = "C00000"
WHITE = "FFFFFF"
LIGHT_BLUE = "D9EAF7"
THIN = Side(style="thin", color="B7B7B7")


def header(ws, row=1):
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.font = Font(name="微软雅黑", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ws.row_dimensions[row].height = 28


def body(ws, start_row=2):
    for row in ws.iter_rows(min_row=start_row):
        for cell in row:
            cell.font = Font(name="微软雅黑", size=10, color="222222")
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def finish(ws, header_row=1, freeze=None, autofilter=True):
    header(ws, header_row)
    body(ws, header_row + 1)
    ws.freeze_panes = freeze or f"A{header_row + 1}"
    if autofilter:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.sheet_view.showGridLines = False
    for col in range(1, ws.max_column + 1):
        vals = [str(ws.cell(r, col).value or "") for r in range(1, min(ws.max_row, 80) + 1)]
        width = min(max(max((len(v) for v in vals), default=8) * 1.2 + 2, 10), 30)
        ws.column_dimensions[get_column_letter(col)].width = width


def table(ws, name, ref=None):
    ref = ref or f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    obj = Table(displayName=name, ref=ref)
    obj.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(obj)


def money(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = '¥#,##0.00;[Red]-¥#,##0.00'


def percent(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = "0.00%"


def save(wb, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(path)


def solve(input_dir, output_dir):
    source_wb = load_workbook(input_dir / "销售订单与目标.xlsx", data_only=True)
    orders = list(source_wb["订单"].iter_rows(min_row=2, values_only=True))
    targets = {row[0]: row[1] + row[2] for row in source_wb["销售目标"].iter_rows(min_row=2, values_only=True)}
    wb = Workbook()
    detail = wb.active
    detail.title = "明细"
    detail.append(["订单号", "日期", "区域", "产品", "数量", "单价", "单位成本", "销售额", "成本", "毛利"])
    enriched = []
    for row in orders:
        sales = row[4] * row[5]
        cost = row[4] * row[6]
        item = list(row) + [sales, cost, sales - cost]
        enriched.append(item)
        detail.append(item)
    finish(detail)
    money(detail, ["F", "G", "H", "I", "J"])
    table(detail, "SalesDetails")

    monthly = defaultdict(lambda: [0.0, 0.0])
    regional = defaultdict(lambda: [0, 0.0, 0.0])
    product = defaultdict(lambda: [0.0, 0.0])
    for row in enriched:
        month = row[1].strftime("%Y-%m")
        monthly[month][0] += row[7]
        monthly[month][1] += row[9]
        regional[row[2]][0] += 1
        regional[row[2]][1] += row[7]
        regional[row[2]][2] += row[9]
        product[row[3]][0] += row[7]
        product[row[3]][1] += row[9]
    analysis = wb.create_sheet("分析")
    analysis.append(["月份", "销售额", "毛利", "毛利率", None, "区域", "订单数", "销售额", "毛利", "目标", "达成率", None, "产品", "销售额", "毛利", "毛利贡献率"])
    max_rows = max(len(monthly), len(regional), len(product))
    months, regions, products = sorted(monthly), sorted(regional), sorted(product)
    total_gross = sum(values[1] for values in product.values())
    for idx in range(max_rows):
        row = [None] * 16
        if idx < len(months):
            month = months[idx]
            row[0:4] = [month, monthly[month][0], monthly[month][1], monthly[month][1] / monthly[month][0]]
        if idx < len(regions):
            region = regions[idx]
            row[5:11] = [region, regional[region][0], regional[region][1], regional[region][2], targets[region], regional[region][1] / targets[region]]
        if idx < len(products):
            name = products[idx]
            row[12:16] = [name, product[name][0], product[name][1], product[name][1] / total_gross]
        analysis.append(row)
    header(analysis, 1)
    body(analysis, 2)
    analysis.freeze_panes = "A2"
    analysis.sheet_view.showGridLines = False
    for col in ["B", "C", "H", "I", "J", "N", "O"]:
        money(analysis, [col], 2, analysis.max_row)
    for col in ["D", "K", "P"]:
        percent(analysis, [col], 2, analysis.max_row)
    for col in range(1, 17):
        analysis.column_dimensions[get_column_letter(col)].width = 14

    dashboard = wb.create_sheet("仪表盘")
    dashboard.append(["2026 年销售运营仪表盘", None, None, None, None])
    dashboard.merge_cells("A1:E1")
    dashboard["A1"].font = Font(size=16, bold=True, color=WHITE)
    dashboard["A1"].fill = PatternFill("solid", fgColor=NAVY)
    dashboard["A1"].alignment = Alignment(horizontal="center")
    total_sales = sum(row[7] for row in enriched)
    total_cost = sum(row[8] for row in enriched)
    kpis = [
        ("A3", "总销售额", total_sales), ("B3", "总毛利", total_sales - total_cost),
        ("C3", "整体毛利率", (total_sales - total_cost) / total_sales), ("D3", "订单数", len(enriched)),
        ("E3", "目标达成率", total_sales / sum(targets.values())),
    ]
    for cell_ref, label, value in kpis:
        cell = dashboard[cell_ref]
        cell.value = f"{label}\n{value:,.2f}" if "率" not in label else f"{label}\n{value:.2%}"
        cell.font = Font(size=12, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    dashboard.row_dimensions[3].height = 50
    for col in range(1, 6):
        dashboard.column_dimensions[get_column_letter(col)].width = 22
    dashboard.sheet_view.showGridLines = False

    line = LineChart()
    line.title = "月度销售趋势"
    line.add_data(Reference(analysis, min_col=2, min_row=1, max_row=len(months) + 1), titles_from_data=True)
    line.set_categories(Reference(analysis, min_col=1, min_row=2, max_row=len(months) + 1))
    line.height, line.width = 7, 13
    dashboard.add_chart(line, "A6")
    region_chart = BarChart()
    region_chart.title = "区域目标达成率"
    region_chart.add_data(Reference(analysis, min_col=11, min_row=1, max_row=len(regions) + 1), titles_from_data=True)
    region_chart.set_categories(Reference(analysis, min_col=6, min_row=2, max_row=len(regions) + 1))
    region_chart.y_axis.numFmt = "0%"
    region_chart.height, region_chart.width = 7, 13
    dashboard.add_chart(region_chart, "H6")
    pie = PieChart()
    pie.title = "产品毛利贡献"
    pie.add_data(Reference(analysis, min_col=15, min_row=1, max_row=len(products) + 1), titles_from_data=True)
    pie.set_categories(Reference(analysis, min_col=13, min_row=2, max_row=len(products) + 1))
    pie.height, pie.width = 7, 13
    dashboard.add_chart(pie, "A21")
    dashboard.sheet_properties.pageSetUpPr.fitToPage = True
    dashboard.page_setup.orientation = "landscape"
    dashboard.page_setup.paperSize = dashboard.PAPERSIZE_A4
    dashboard.page_setup.fitToWidth = 1
    dashboard.page_setup.fitToHeight = 1
    dashboard.print_area = "A1:T38"
    dashboard.page_margins.left = 0.25
    dashboard.page_margins.right = 0.25
    dashboard.page_margins.top = 0.35
    dashboard.page_margins.bottom = 0.35
    save(wb, output_dir / "销售运营仪表盘.xlsx")



def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--output-dir", required=True)
    args = parser.parse_args()
    input_dir = Path(args.input_dir).resolve()
    output_dir = Path(args.output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    solve(input_dir, output_dir)


if __name__ == "__main__":
    main()
