
from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "17365D"
BLUE = "4472C4"
GREEN = "70AD47"
ORANGE = "ED7D31"
RED = "C00000"
WHITE = "FFFFFF"
LIGHT_BLUE = "D9EAF7"
THIN = Side(style="thin", color="B7B7B7")


def header(ws, row=1):
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.font = Font(name="微软雅黑", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ws.row_dimensions[row].height = 28


def body(ws, start_row=2):
    for row in ws.iter_rows(min_row=start_row):
        for cell in row:
            cell.font = Font(name="微软雅黑", size=10, color="222222")
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def finish(ws, header_row=1, freeze=None, autofilter=True):
    header(ws, header_row)
    body(ws, header_row + 1)
    ws.freeze_panes = freeze or f"A{header_row + 1}"
    if autofilter:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.sheet_view.showGridLines = False
    for col in range(1, ws.max_column + 1):
        vals = [str(ws.cell(r, col).value or "") for r in range(1, min(ws.max_row, 80) + 1)]
        width = min(max(max((len(v) for v in vals), default=8) * 1.2 + 2, 10), 30)
        ws.column_dimensions[get_column_letter(col)].width = width


def table(ws, name, ref=None):
    ref = ref or f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    obj = Table(displayName=name, ref=ref)
    obj.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(obj)


def money(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = '¥#,##0.00;[Red]-¥#,##0.00'


def percent(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = "0.00%"


def save(wb, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(path)


def parse_budget(raw):
    text = str(raw).strip().replace("¥", "").replace(",", "")
    if not text or text.upper() in {"N/A", "NA", "NONE"}:
        return None
    multiplier = 10000 if text.endswith("万") else 1
    text = text[:-1] if text.endswith("万") else text
    try:
        return round(float(text) * multiplier, 2)
    except ValueError:
        return None


def solve(input_dir, output_dir):
    mapping_wb = load_workbook(input_dir / "省份映射.xlsx", data_only=True)
    mapping = {str(row[0]).strip(): str(row[1]).strip() for row in mapping_wb.active.iter_rows(min_row=2, values_only=True)}
    records = {}
    input_count = 0
    with (input_dir / "客户线索_待清洗.csv").open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            input_count += 1
            key = row["lead_id"].strip()
            if key not in records or row["updated_at"] > records[key]["updated_at"]:
                records[key] = row

    cleaned = []
    for key in sorted(records):
        row = records[key]
        name = row["name"].strip()
        phone = re.sub(r"\D", "", row["phone"])
        if len(phone) == 13 and phone.startswith("86"):
            phone = phone[2:]
        email = re.sub(r"\s+", "", row["email"]).lower()
        province = mapping.get(row["province"].strip(), "未知")
        budget = parse_budget(row["budget"])
        consent = "是" if row["consent"].strip().lower() in {"是", "y", "yes", "true", "1"} else "否"
        problems = []
        if not re.fullmatch(r"1[3-9]\d{9}", phone):
            problems.append("手机号无效")
        if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
            problems.append("邮箱无效")
        if province == "未知":
            problems.append("省份未匹配")
        if budget is None or budget < 0:
            problems.append("预算无效")
        cleaned.append([
            key, name, phone, email, province, row["created_at"].replace("/", "-").replace(".", "-"),
            budget, consent, row["updated_at"], "通过" if not problems else "；".join(problems),
        ])

    wb = Workbook()
    ws = wb.active
    ws.title = "清洗结果"
    ws.append(["lead_id", "姓名", "手机号", "邮箱", "标准省份", "创建日期", "预算（元）", "同意营销", "最后更新时间", "校验状态"])
    for row in cleaned:
        ws.append(row)
    finish(ws)
    money(ws, ["G"])
    table(ws, "CleanLeads")
    ws.conditional_formatting.add(
        f"J2:J{ws.max_row}", CellIsRule(operator="notEqual", formula=['"通过"'], fill=PatternFill("solid", fgColor="F4CCCC"))
    )

    summary = wb.create_sheet("质量汇总")
    pass_count = sum(1 for row in cleaned if row[-1] == "通过")
    summary.append(["指标", "数值"])
    for row in [
        ["输入记录数", input_count], ["去重后记录数", len(cleaned)], ["重复记录数", input_count - len(cleaned)],
        ["校验通过数", pass_count], ["校验不通过数", len(cleaned) - pass_count], ["通过率", pass_count / len(cleaned)],
    ]:
        summary.append(row)
    finish(summary)
    summary["B7"].number_format = "0.00%"

    province_ws = wb.create_sheet("省份统计")
    counts = defaultdict(int)
    for row in cleaned:
        counts[row[4]] += 1
    province_ws.append(["标准省份", "线索数"])
    for province in sorted(counts):
        province_ws.append([province, counts[province]])
    finish(province_ws)
    save(wb, output_dir / "客户线索清洗结果.xlsx")



def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--output-dir", required=True)
    args = parser.parse_args()
    input_dir = Path(args.input_dir).resolve()
    output_dir = Path(args.output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    solve(input_dir, output_dir)


if __name__ == "__main__":
    main()
