
from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "17365D"
BLUE = "4472C4"
GREEN = "70AD47"
ORANGE = "ED7D31"
RED = "C00000"
WHITE = "FFFFFF"
LIGHT_BLUE = "D9EAF7"
THIN = Side(style="thin", color="B7B7B7")


def header(ws, row=1):
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.font = Font(name="微软雅黑", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ws.row_dimensions[row].height = 28


def body(ws, start_row=2):
    for row in ws.iter_rows(min_row=start_row):
        for cell in row:
            cell.font = Font(name="微软雅黑", size=10, color="222222")
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def finish(ws, header_row=1, freeze=None, autofilter=True):
    header(ws, header_row)
    body(ws, header_row + 1)
    ws.freeze_panes = freeze or f"A{header_row + 1}"
    if autofilter:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.sheet_view.showGridLines = False
    for col in range(1, ws.max_column + 1):
        vals = [str(ws.cell(r, col).value or "") for r in range(1, min(ws.max_row, 80) + 1)]
        width = min(max(max((len(v) for v in vals), default=8) * 1.2 + 2, 10), 30)
        ws.column_dimensions[get_column_letter(col)].width = width


def table(ws, name, ref=None):
    ref = ref or f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    obj = Table(displayName=name, ref=ref)
    obj.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(obj)


def money(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = '¥#,##0.00;[Red]-¥#,##0.00'


def percent(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = "0.00%"


def save(wb, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(path)


def solve(input_dir, output_dir):
    document = Document(input_dir / "供应商报价纪要.docx")
    records = []
    for paragraph in document.paragraphs:
        parts = [part.strip() for part in paragraph.text.split("|")]
        if len(parts) == 7 and parts[0].isdigit():
            records.append({
                "seq": int(parts[0]), "vendor": parts[1], "code": parts[2], "name": parts[3],
                "quantity": int(parts[4]), "unit_price": float(parts[5]), "days": int(parts[6]),
            })
    if len(records) != 6:
        raise ValueError(f"期望 6 条报价，实际读取 {len(records)} 条")
    for item in records:
        item["total"] = round(item["quantity"] * item["unit_price"], 2)
    records.sort(key=lambda item: (item["code"], item["total"], item["seq"]))

    wb = Workbook()
    detail = wb.active
    detail.title = "报价明细"
    detail.append(["报价序号", "供应商", "物料编码", "物料名称", "数量", "含税单价", "含税总额", "交付天数"])
    for item in records:
        detail.append([item["seq"], item["vendor"], item["code"], item["name"], item["quantity"], item["unit_price"], item["total"], item["days"]])
    finish(detail)
    money(detail, ["F", "G"])
    table(detail, "QuoteDetails")

    vendor_summary = defaultdict(lambda: [0, 0.0, []])
    for item in records:
        summary = vendor_summary[item["vendor"]]
        summary[0] += 1
        summary[1] += item["total"]
        summary[2].append(item["days"])
    summary_ws = wb.create_sheet("供应商汇总")
    summary_ws.append(["供应商", "报价条数", "报价金额合计", "平均交付天数"])
    for vendor in sorted(vendor_summary):
        count, total, days = vendor_summary[vendor]
        summary_ws.append([vendor, count, round(total, 2), round(sum(days) / len(days), 2)])
    finish(summary_ws)
    money(summary_ws, ["C"])
    table(summary_ws, "VendorSummary")

    best_by_code = {}
    for item in records:
        best_by_code.setdefault(item["code"], item)
    recommendation = wb.create_sheet("最低价推荐")
    recommendation.append(["物料编码", "物料名称", "推荐供应商", "数量", "含税单价", "含税总额", "交付天数"])
    for code in sorted(best_by_code):
        item = best_by_code[code]
        recommendation.append([code, item["name"], item["vendor"], item["quantity"], item["unit_price"], item["total"], item["days"]])
    recommendation.append(["推荐合计", None, None, None, None, sum(item["total"] for item in best_by_code.values()), None])
    finish(recommendation)
    money(recommendation, ["E", "F"])
    recommendation[f"A{recommendation.max_row}"].font = Font(bold=True, color=WHITE)
    recommendation[f"A{recommendation.max_row}"].fill = PatternFill("solid", fgColor=NAVY)
    recommendation[f"F{recommendation.max_row}"].font = Font(bold=True)
    save(wb, output_dir / "采购报价结构化.xlsx")



def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--output-dir", required=True)
    args = parser.parse_args()
    input_dir = Path(args.input_dir).resolve()
    output_dir = Path(args.output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    solve(input_dir, output_dir)


if __name__ == "__main__":
    main()
