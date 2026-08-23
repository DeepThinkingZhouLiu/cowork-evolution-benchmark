#!/usr/bin/env python3
"""确定性 Office 任务 Judge。

Golden 和 checks 属于评测侧隐藏资产，不得挂载到 Agent 可见的输入目录。
"""
from __future__ import annotations

import argparse
import json
import math
import re
import sys
from datetime import date, datetime, time
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

ERROR_MARKERS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#NULL!", "#N/A", "#NUM!"}


def normalize_formula(value):
    if not isinstance(value, str) or not value.startswith("="):
        return value
    return re.sub(r"\s+", "", value.replace("$", "").upper())


def normalize(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def equivalent(actual, expected, tolerance):
    if isinstance(actual, str) and actual.startswith("=") and isinstance(expected, str) and expected.startswith("="):
        return normalize_formula(actual) == normalize_formula(expected)
    actual, expected = normalize(actual), normalize(expected)
    if actual in (None, "") and expected in (None, ""):
        return True
    if isinstance(actual, bool) or isinstance(expected, bool):
        return actual is expected
    if isinstance(actual, (int, float)) and isinstance(expected, (int, float)):
        if math.isnan(actual) or math.isnan(expected):
            return False
        return abs(float(actual) - float(expected)) <= tolerance
    return actual == expected


def color_rgb(color):
    if color is None:
        return None
    if getattr(color, "type", None) == "rgb" and isinstance(color.rgb, str):
        return color.rgb[-6:].upper()
    if getattr(color, "type", None) == "theme":
        return f"THEME:{color.theme}"
    return None


def compare_range(actual_ws, golden_ws, range_ref, tolerance, errors):
    for row in golden_ws[range_ref]:
        for golden_cell in row:
            actual_cell = actual_ws[golden_cell.coordinate]
            if not equivalent(actual_cell.value, golden_cell.value, tolerance):
                errors.append(
                    f"{actual_ws.title}!{golden_cell.coordinate}: "
                    f"期望 {golden_cell.value!r}，实际 {actual_cell.value!r}"
                )


def chart_type(chart):
    return type(chart).__name__


def scan_errors(wb):
    found = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if value in ERROR_MARKERS or (
                    isinstance(value, str) and value.startswith("=") and any(marker in value for marker in ERROR_MARKERS)
                ):
                    found.append(f"{ws.title}!{cell.coordinate}={value}")
    return found


def check_xlsx(candidate, golden, config, tolerance):
    result = {"integrity": [], "structure": [], "accuracy": [], "formula": [], "format": [], "quality": []}
    try:
        actual_wb = load_workbook(candidate, data_only=False)
        golden_wb = load_workbook(golden, data_only=False)
        result["integrity"].append("工作簿可正常打开")
    except Exception as exc:
        result["integrity"].append(f"ERROR:工作簿无法打开: {exc}")
        return result

    required = config.get("sheets", [])
    for name in required:
        if name not in actual_wb.sheetnames:
            result["structure"].append(f"ERROR:缺少工作表 {name}")
        else:
            result["structure"].append(f"工作表 {name} 存在")

    if config.get("exact_sheet_order") and actual_wb.sheetnames[: len(required)] != required:
        result["structure"].append(
            f"ERROR:工作表顺序不正确: {actual_wb.sheetnames}"
        )

    for item in config.get("ranges", []):
        sheet = item["sheet"]
        if sheet not in actual_wb.sheetnames or sheet not in golden_wb.sheetnames:
            continue
        errors = []
        compare_range(actual_wb[sheet], golden_wb[sheet], item["range"], tolerance, errors)
        if errors:
            result["accuracy"].extend("ERROR:" + msg for msg in errors[:40])
            if len(errors) > 40:
                result["accuracy"].append(f"ERROR:还有 {len(errors) - 40} 个单元格不匹配")
        else:
            result["accuracy"].append(f"{sheet}!{item['range']} 与标准结果一致")

        formula_total = 0
        formula_bad = 0
        for row in golden_wb[sheet][item["range"]]:
            for golden_cell in row:
                if isinstance(golden_cell.value, str) and golden_cell.value.startswith("="):
                    formula_total += 1
                    actual_value = actual_wb[sheet][golden_cell.coordinate].value
                    if normalize_formula(actual_value) != normalize_formula(golden_cell.value):
                        formula_bad += 1
        if formula_total:
            if formula_bad:
                result["formula"].append(
                    f"ERROR:{sheet}!{item['range']} 有 {formula_bad}/{formula_total} 个公式不正确"
                )
            else:
                result["formula"].append(f"{sheet} 的 {formula_total} 个关键公式正确")

    for sheet, rows in config.get("header_rows", {}).items():
        if sheet not in actual_wb.sheetnames:
            continue
        actual_ws, golden_ws = actual_wb[sheet], golden_wb[sheet]
        for row_idx in rows:
            for col_idx in range(1, golden_ws.max_column + 1):
                golden_cell = golden_ws.cell(row_idx, col_idx)
                if golden_cell.value is None:
                    continue
                actual_cell = actual_ws.cell(row_idx, col_idx)
                if actual_cell.font.bold != golden_cell.font.bold:
                    result["format"].append(
                        f"ERROR:{sheet}!{actual_cell.coordinate} 表头粗体样式不正确"
                    )
                if color_rgb(actual_cell.fill.fgColor) != color_rgb(golden_cell.fill.fgColor):
                    result["format"].append(
                        f"ERROR:{sheet}!{actual_cell.coordinate} 表头填充色不正确"
                    )
        if not any(x.startswith("ERROR:") and x.startswith(f"ERROR:{sheet}!") for x in result["format"]):
            result["format"].append(f"{sheet} 表头样式正确")

    for sheet, expected in config.get("freeze_panes", {}).items():
        if sheet in actual_wb.sheetnames:
            value = actual_wb[sheet].freeze_panes
            value = value.coordinate if hasattr(value, "coordinate") else value
            if value != expected:
                result["format"].append(f"ERROR:{sheet} 冻结窗格应为 {expected}，实际为 {value}")

    for sheet, expected_count in config.get("min_charts", {}).items():
        if sheet in actual_wb.sheetnames:
            actual_charts = actual_wb[sheet]._charts
            golden_charts = golden_wb[sheet]._charts
            if len(actual_charts) < expected_count:
                result["format"].append(
                    f"ERROR:{sheet} 至少需要 {expected_count} 张可编辑图表，实际 {len(actual_charts)}"
                )
            elif [chart_type(c) for c in actual_charts[:expected_count]] != [
                chart_type(c) for c in golden_charts[:expected_count]
            ]:
                result["format"].append(f"ERROR:{sheet} 图表类型与要求不一致")
            else:
                result["format"].append(f"{sheet} 包含 {len(actual_charts)} 张类型正确的可编辑图表")

    for sheet, state in config.get("sheet_states", {}).items():
        if sheet in actual_wb.sheetnames and actual_wb[sheet].sheet_state != state:
            result["format"].append(f"ERROR:{sheet} 的工作表状态应为 {state}")

    errors = scan_errors(actual_wb)
    if errors:
        result["quality"].extend("ERROR:检测到表格错误 " + item for item in errors[:20])
    else:
        result["quality"].append("未检测到 Excel 错误标记")

    nonempty = sum(
        1 for ws in actual_wb.worksheets for row in ws.iter_rows() for cell in row if cell.value not in (None, "")
    )
    if nonempty < config.get("min_nonempty_cells", 10):
        result["quality"].append(f"ERROR:工作簿内容过少，仅 {nonempty} 个非空单元格")
    return result


def check_docx(candidate, golden, config):
    result = {"integrity": [], "structure": [], "accuracy": [], "formula": [], "format": [], "quality": []}
    try:
        actual = Document(candidate)
        expected = Document(golden)
        result["integrity"].append("Word 报告可正常打开")
    except Exception as exc:
        result["integrity"].append(f"ERROR:Word 报告无法打开: {exc}")
        return result

    actual_text = "\n".join(p.text for p in actual.paragraphs)
    expected_text = "\n".join(p.text for p in expected.paragraphs)
    for token in config.get("required_text", []):
        if token not in actual_text:
            result["structure"].append(f"ERROR:Word 报告缺少文本: {token}")
        else:
            result["structure"].append(f"Word 报告包含: {token}")

    for token in config.get("golden_numeric_tokens", []):
        if token not in actual_text or token not in expected_text:
            result["accuracy"].append(f"ERROR:Word 报告缺少关键数字: {token}")
        else:
            result["accuracy"].append(f"Word 报告关键数字正确: {token}")

    heading_count = sum(1 for p in actual.paragraphs if p.style and p.style.name.startswith("Heading"))
    if heading_count < config.get("min_headings", 3):
        result["format"].append(f"ERROR:Word 报告标题层级不完整，仅 {heading_count} 个标题")
    else:
        result["format"].append(f"Word 报告包含 {heading_count} 个标题")

    if len(actual.paragraphs) < config.get("min_paragraphs", 10):
        result["quality"].append("ERROR:Word 报告段落过少")
    else:
        result["quality"].append(f"Word 报告包含 {len(actual.paragraphs)} 个段落")
    if len(actual.tables) < config.get("min_tables", 1):
        result["structure"].append("ERROR:Word 报告缺少 KPI 表格")
    return result


def has_error(messages):
    return any(item.startswith("ERROR:") or ": ERROR:" in item for item in messages)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--candidate-dir", required=True)
    parser.add_argument("--golden-dir", required=True)
    parser.add_argument("--checks", required=True)
    parser.add_argument("--result-json")
    args = parser.parse_args()

    candidate_dir = Path(args.candidate_dir)
    golden_dir = Path(args.golden_dir)
    checks = json.loads(Path(args.checks).read_text(encoding="utf-8"))
    grouped = {key: [] for key in ["integrity", "structure", "accuracy", "formula", "format", "quality"]}

    for file_config in checks["files"]:
        name = file_config["name"]
        candidate = candidate_dir / name
        golden = golden_dir / name
        if not candidate.is_file():
            grouped["integrity"].append(f"ERROR:缺少必需文件 {name}")
            continue
        if candidate.stat().st_size < 500:
            grouped["integrity"].append(f"ERROR:{name} 文件过小，疑似空壳")
            continue
        if file_config["type"] == "xlsx":
            partial = check_xlsx(candidate, golden, file_config, checks.get("numeric_tolerance", 0.01))
        else:
            partial = check_docx(candidate, golden, file_config)
        for key, items in partial.items():
            grouped[key].extend(f"{name}: {item}" for item in items)

    criteria = [
        {"id": "R001", "group": "integrity", "weight": 2},
        {"id": "R002", "group": "structure", "weight": 2},
        {"id": "R003", "group": "accuracy", "weight": 4},
        {"id": "R004", "group": "formula", "weight": 3},
        {"id": "R005", "group": "format", "weight": 2},
        {"id": "R006", "group": "quality", "weight": 1},
    ]
    scored = []
    earned = 0
    total = sum(item["weight"] for item in criteria)
    for item in criteria:
        messages = grouped[item["group"]]
        passed = bool(messages) and not has_error(messages)
        if item["group"] == "formula" and not messages:
            # 没有指定公式的任务不因“无公式”失分。
            passed = True
            messages = ["本任务无额外强制公式检查"]
        if passed:
            earned += item["weight"]
        scored.append(
            {
                "id": item["id"],
                "passed": passed,
                "weight": item["weight"],
                "evidence": messages[:12],
            }
        )

    penalty = has_error(grouped["quality"]) or has_error(grouped["integrity"])
    if penalty:
        earned = max(0, earned - 4)
    score = earned / total if total else 0
    hurdle_passed = scored[0]["passed"]
    passed = hurdle_passed and score >= checks.get("pass_threshold", 0.8)
    result = {
        "task_id": checks["task_id"],
        "passed": passed,
        "score": round(score, 6),
        "pass_threshold": checks.get("pass_threshold", 0.8),
        "criteria": scored,
        "penalty": {"id": "R007", "triggered": penalty, "weight": -4},
    }
    payload = json.dumps(result, ensure_ascii=False, indent=2)
    print(payload)
    if args.result_json:
        Path(args.result_json).write_text(payload + "\n", encoding="utf-8")
    raise SystemExit(0 if passed else 1)


if __name__ == "__main__":
    main()
