
from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "17365D"
BLUE = "4472C4"
GREEN = "70AD47"
ORANGE = "ED7D31"
RED = "C00000"
WHITE = "FFFFFF"
LIGHT_BLUE = "D9EAF7"
THIN = Side(style="thin", color="B7B7B7")


def header(ws, row=1):
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.font = Font(name="微软雅黑", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ws.row_dimensions[row].height = 28


def body(ws, start_row=2):
    for row in ws.iter_rows(min_row=start_row):
        for cell in row:
            cell.font = Font(name="微软雅黑", size=10, color="222222")
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def finish(ws, header_row=1, freeze=None, autofilter=True):
    header(ws, header_row)
    body(ws, header_row + 1)
    ws.freeze_panes = freeze or f"A{header_row + 1}"
    if autofilter:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.sheet_view.showGridLines = False
    for col in range(1, ws.max_column + 1):
        vals = [str(ws.cell(r, col).value or "") for r in range(1, min(ws.max_row, 80) + 1)]
        width = min(max(max((len(v) for v in vals), default=8) * 1.2 + 2, 10), 30)
        ws.column_dimensions[get_column_letter(col)].width = width


def table(ws, name, ref=None):
    ref = ref or f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    obj = Table(displayName=name, ref=ref)
    obj.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(obj)


def money(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = '¥#,##0.00;[Red]-¥#,##0.00'


def percent(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = "0.00%"


def save(wb, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(path)


def solve(input_dir, output_dir):
    wb = load_workbook(input_dir / "SaaS经营假设与历史.xlsx", data_only=False)
    model = wb.create_sheet("三年模型")
    model.merge_cells("A1:D1")
    model["A1"] = "SaaS 2026–2028 三年经营模型"
    model["A1"].font = Font(name="微软雅黑", size=15, bold=True, color=WHITE)
    model["A1"].fill = PatternFill("solid", fgColor=NAVY)
    model["A1"].alignment = Alignment(horizontal="center")
    model.append(["指标", 2026, 2027, 2028])
    metrics = ["期初客户数", "新增客户", "流失客户", "期末客户数", "平均客户数", "月 ARPU", "订阅收入", "销售成本率", "销售成本", "毛利", "毛利率", "薪酬", "营销费用", "其他费用", "EBITDA", "EBITDA 率"]
    for metric in metrics:
        model.append([metric, None, None, None])
    for col, letter in enumerate(["B", "C", "D"], 2):
        assumption_col = get_column_letter(col)
        model[f"{letter}3"] = "='历史数据'!C2" if letter == "B" else f"={get_column_letter(col-1)}6"
        model[f"{letter}4"] = f"='经营假设'!{assumption_col}2"
        model[f"{letter}5"] = f"={letter}3*'经营假设'!{assumption_col}3"
        model[f"{letter}6"] = f"={letter}3+{letter}4-{letter}5"
        model[f"{letter}7"] = f"=AVERAGE({letter}3,{letter}6)"
        model[f"{letter}8"] = f"='经营假设'!{assumption_col}4"
        model[f"{letter}9"] = f"={letter}7*{letter}8*12"
        model[f"{letter}10"] = f"='经营假设'!{assumption_col}5"
        model[f"{letter}11"] = f"={letter}9*{letter}10"
        model[f"{letter}12"] = f"={letter}9-{letter}11"
        model[f"{letter}13"] = f"=IFERROR({letter}12/{letter}9,0)"
        model[f"{letter}14"] = f"='经营假设'!{assumption_col}6"
        model[f"{letter}15"] = f"='经营假设'!{assumption_col}7"
        model[f"{letter}16"] = f"='经营假设'!{assumption_col}8"
        model[f"{letter}17"] = f"={letter}12-SUM({letter}14:{letter}16)"
        model[f"{letter}18"] = f"=IFERROR({letter}17/{letter}9,0)"
    finish(model, header_row=2, freeze="B3", autofilter=False)
    for row in [8, 9, 11, 12, 14, 15, 16, 17]:
        money(model, ["B", "C", "D"], row, row)
    for row in [10, 13, 18]:
        percent(model, ["B", "C", "D"], row, row)
    for row in [4, 5, 8, 10, 14, 15, 16]:
        for col in range(2, 5):
            model.cell(row, col).fill = PatternFill("solid", fgColor=LIGHT_BLUE)

    summary = wb.create_sheet("指标摘要")
    summary.append(["指标", "数值"])
    summary.append(["收入 CAGR（2026–2028）", "=IFERROR(('三年模型'!D9/'三年模型'!B9)^(1/2)-1,0)"])
    summary.append(["2028 年毛利率", "='三年模型'!D13"])
    summary.append(["2028 年 EBITDA 率", "='三年模型'!D18"])
    finish(summary)
    for row in range(2, 5):
        summary[f"B{row}"].number_format = "0.00%"
    save(wb, output_dir / "SaaS三年经营模型.xlsx")



def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--output-dir", required=True)
    args = parser.parse_args()
    input_dir = Path(args.input_dir).resolve()
    output_dir = Path(args.output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    solve(input_dir, output_dir)


if __name__ == "__main__":
    main()
