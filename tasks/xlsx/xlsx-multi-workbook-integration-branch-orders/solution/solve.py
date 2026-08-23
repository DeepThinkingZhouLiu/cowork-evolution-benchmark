
from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "17365D"
BLUE = "4472C4"
GREEN = "70AD47"
ORANGE = "ED7D31"
RED = "C00000"
WHITE = "FFFFFF"
LIGHT_BLUE = "D9EAF7"
THIN = Side(style="thin", color="B7B7B7")


def header(ws, row=1):
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.font = Font(name="微软雅黑", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ws.row_dimensions[row].height = 28


def body(ws, start_row=2):
    for row in ws.iter_rows(min_row=start_row):
        for cell in row:
            cell.font = Font(name="微软雅黑", size=10, color="222222")
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def finish(ws, header_row=1, freeze=None, autofilter=True):
    header(ws, header_row)
    body(ws, header_row + 1)
    ws.freeze_panes = freeze or f"A{header_row + 1}"
    if autofilter:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.sheet_view.showGridLines = False
    for col in range(1, ws.max_column + 1):
        vals = [str(ws.cell(r, col).value or "") for r in range(1, min(ws.max_row, 80) + 1)]
        width = min(max(max((len(v) for v in vals), default=8) * 1.2 + 2, 10), 30)
        ws.column_dimensions[get_column_letter(col)].width = width


def table(ws, name, ref=None):
    ref = ref or f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    obj = Table(displayName=name, ref=ref)
    obj.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(obj)


def money(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = '¥#,##0.00;[Red]-¥#,##0.00'


def percent(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = "0.00%"


def save(wb, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(path)


def solve(input_dir, output_dir):
    master_wb = load_workbook(input_dir / "产品主数据.xlsx", data_only=True)
    master = {row[0]: (row[1], row[2]) for row in master_wb.active.iter_rows(min_row=2, values_only=True)}
    selected = {}
    discarded = []
    for region in ["华东", "华南", "华北"]:
        wb = load_workbook(input_dir / f"{region}订单.xlsx", data_only=True)
        for row in wb.active.iter_rows(min_row=2, values_only=True):
            record = {"order": row[0], "date": row[1], "sku": row[2], "quantity": row[3], "price": row[4], "updated": row[5], "region": region}
            previous = selected.get(record["order"])
            if previous is None or record["updated"] > previous["updated"]:
                if previous is not None:
                    discarded.append([previous["order"], previous["region"], previous["updated"], region, record["updated"], "保留更新时间最晚的记录"])
                selected[record["order"]] = record
            else:
                discarded.append([record["order"], region, record["updated"], previous["region"], previous["updated"], "保留更新时间最晚的记录"])

    wb = Workbook()
    detail = wb.active
    detail.title = "合并明细"
    detail.append(["订单号", "订单日期", "来源分公司", "SKU", "品类", "数量", "成交单价", "销售额", "标准单位成本", "标准成本", "标准毛利", "更新时间", "主数据状态"])
    enriched = []
    for order in sorted(selected):
        item = selected[order]
        category, unit_cost = master.get(item["sku"], ("未匹配", None))
        sales = item["quantity"] * item["price"]
        cost = item["quantity"] * unit_cost if unit_cost is not None else None
        gross = sales - cost if cost is not None else None
        row = [order, item["date"], item["region"], item["sku"], category, item["quantity"], item["price"], sales, unit_cost, cost, gross, item["updated"], "已匹配" if unit_cost is not None else "未匹配"]
        enriched.append(row)
        detail.append(row)
    finish(detail)
    money(detail, ["G", "H", "I", "J", "K"])
    detail.conditional_formatting.add(f"M2:M{detail.max_row}", CellIsRule(operator="equal", formula=['"未匹配"'], fill=PatternFill("solid", fgColor="F4CCCC")))
    table(detail, "IntegratedOrders")

    summary = wb.create_sheet("分公司汇总")
    grouped = defaultdict(lambda: [0, 0.0, 0.0])
    for row in enriched:
        grouped[row[2]][0] += 1
        grouped[row[2]][1] += row[7]
        if row[10] is not None:
            grouped[row[2]][2] += row[10]
    summary.append(["来源分公司", "订单数", "销售额", "标准毛利（已匹配 SKU）"])
    for region in sorted(grouped):
        summary.append([region] + grouped[region])
    finish(summary)
    money(summary, ["C", "D"])

    reconciliation = wb.create_sheet("对账清单")
    reconciliation.append(["订单号", "被舍弃来源", "被舍弃更新时间", "保留来源", "保留更新时间", "原因"])
    for row in sorted(discarded):
        reconciliation.append(row)
    finish(reconciliation)
    save(wb, output_dir / "分公司订单汇总.xlsx")



def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--output-dir", required=True)
    args = parser.parse_args()
    input_dir = Path(args.input_dir).resolve()
    output_dir = Path(args.output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    solve(input_dir, output_dir)


if __name__ == "__main__":
    main()
