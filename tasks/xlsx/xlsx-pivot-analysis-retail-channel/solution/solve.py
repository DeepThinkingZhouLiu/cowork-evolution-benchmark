
from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "17365D"
BLUE = "4472C4"
GREEN = "70AD47"
ORANGE = "ED7D31"
RED = "C00000"
WHITE = "FFFFFF"
LIGHT_BLUE = "D9EAF7"
THIN = Side(style="thin", color="B7B7B7")


def header(ws, row=1):
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.font = Font(name="微软雅黑", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ws.row_dimensions[row].height = 28


def body(ws, start_row=2):
    for row in ws.iter_rows(min_row=start_row):
        for cell in row:
            cell.font = Font(name="微软雅黑", size=10, color="222222")
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def finish(ws, header_row=1, freeze=None, autofilter=True):
    header(ws, header_row)
    body(ws, header_row + 1)
    ws.freeze_panes = freeze or f"A{header_row + 1}"
    if autofilter:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.sheet_view.showGridLines = False
    for col in range(1, ws.max_column + 1):
        vals = [str(ws.cell(r, col).value or "") for r in range(1, min(ws.max_row, 80) + 1)]
        width = min(max(max((len(v) for v in vals), default=8) * 1.2 + 2, 10), 30)
        ws.column_dimensions[get_column_letter(col)].width = width


def table(ws, name, ref=None):
    ref = ref or f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    obj = Table(displayName=name, ref=ref)
    obj.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(obj)


def money(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = '¥#,##0.00;[Red]-¥#,##0.00'


def percent(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = "0.00%"


def save(wb, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(path)


def solve(input_dir, output_dir):
    source_wb = load_workbook(input_dir / "零售交易明细.xlsx", data_only=True)
    source = source_wb.active
    rows = list(source.iter_rows(min_row=2, values_only=True))
    wb = Workbook()
    detail = wb.active
    detail.title = "交易明细"
    detail.append([cell.value for cell in source[1]] + ["毛利"])
    for row in rows:
        detail.append(list(row) + [row[6] - row[7]])
    finish(detail)
    money(detail, ["G", "H", "I"])
    table(detail, "RetailTransactions")

    channels = sorted({row[3] for row in rows})
    categories = sorted({row[4] for row in rows})
    revenue = defaultdict(float)
    gross = defaultdict(float)
    for row in rows:
        revenue[(row[3], row[4])] += row[6]
        gross[row[3]] += row[6] - row[7]
    pivot = wb.create_sheet("透视分析")
    pivot.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(categories) + 3)
    pivot["A1"] = "渠道 × 品类销售额透视式汇总"
    pivot["A1"].font = Font(size=14, bold=True, color=WHITE)
    pivot["A1"].fill = PatternFill("solid", fgColor=NAVY)
    pivot["A1"].alignment = Alignment(horizontal="center")
    pivot.append(["渠道"] + categories + ["行合计", "毛利率"])
    for channel in channels:
        values = [revenue[(channel, category)] for category in categories]
        total = sum(values)
        pivot.append([channel] + values + [total, gross[channel] / total])
    column_totals = [sum(revenue[(channel, category)] for channel in channels) for category in categories]
    grand_total = sum(column_totals)
    pivot.append(["列合计"] + column_totals + [grand_total, sum(gross.values()) / grand_total])
    header(pivot, 2)
    body(pivot, 3)
    pivot.freeze_panes = "B3"
    for col in range(2, len(categories) + 3):
        money(pivot, [get_column_letter(col)], 3, pivot.max_row)
    percent(pivot, [get_column_letter(len(categories) + 3)], 3, pivot.max_row)
    pivot.sheet_view.showGridLines = False
    for col in range(1, pivot.max_column + 1):
        pivot.column_dimensions[get_column_letter(col)].width = 16

    trend = wb.create_sheet("月度趋势")
    monthly = defaultdict(lambda: [0.0, 0.0])
    for row in rows:
        month = row[1].strftime("%Y-%m")
        monthly[month][0] += row[6]
        monthly[month][1] += row[6] - row[7]
    trend.append(["月份", "销售额", "毛利", "毛利率"])
    for month in sorted(monthly):
        sales, margin = monthly[month]
        trend.append([month, sales, margin, margin / sales])
    finish(trend)
    money(trend, ["B", "C"])
    percent(trend, ["D"])
    chart = BarChart()
    chart.type = "col"
    chart.title = "月度销售额趋势"
    chart.y_axis.title = "销售额（元）"
    chart.x_axis.title = "月份"
    chart.add_data(Reference(trend, min_col=2, min_row=1, max_row=trend.max_row), titles_from_data=True)
    chart.set_categories(Reference(trend, min_col=1, min_row=2, max_row=trend.max_row))
    chart.height = 8
    chart.width = 15
    trend.add_chart(chart, "F2")
    save(wb, output_dir / "零售渠道透视分析.xlsx")



def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--output-dir", required=True)
    args = parser.parse_args()
    input_dir = Path(args.input_dir).resolve()
    output_dir = Path(args.output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    solve(input_dir, output_dir)


if __name__ == "__main__":
    main()
