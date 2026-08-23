
from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "17365D"
BLUE = "4472C4"
GREEN = "70AD47"
ORANGE = "ED7D31"
RED = "C00000"
WHITE = "FFFFFF"
LIGHT_BLUE = "D9EAF7"
THIN = Side(style="thin", color="B7B7B7")


def header(ws, row=1):
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.font = Font(name="微软雅黑", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ws.row_dimensions[row].height = 28


def body(ws, start_row=2):
    for row in ws.iter_rows(min_row=start_row):
        for cell in row:
            cell.font = Font(name="微软雅黑", size=10, color="222222")
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def finish(ws, header_row=1, freeze=None, autofilter=True):
    header(ws, header_row)
    body(ws, header_row + 1)
    ws.freeze_panes = freeze or f"A{header_row + 1}"
    if autofilter:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.sheet_view.showGridLines = False
    for col in range(1, ws.max_column + 1):
        vals = [str(ws.cell(r, col).value or "") for r in range(1, min(ws.max_row, 80) + 1)]
        width = min(max(max((len(v) for v in vals), default=8) * 1.2 + 2, 10), 30)
        ws.column_dimensions[get_column_letter(col)].width = width


def table(ws, name, ref=None):
    ref = ref or f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    obj = Table(displayName=name, ref=ref)
    obj.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(obj)


def money(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = '¥#,##0.00;[Red]-¥#,##0.00'


def percent(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = "0.00%"


def save(wb, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(path)


def solve(input_dir, output_dir):
    wb = load_workbook(input_dir / "物流定价假设.xlsx", data_only=False)
    scenario = wb.create_sheet("情景分析")
    scenario.append(["情景", "调价幅度", "调整后价格", "预计包裹量", "收入", "变动成本", "贡献毛利", "营业利润"])
    for row in range(2, 5):
        scenario[f"A{row}"] = f"='情景参数'!A{row}"
        scenario[f"B{row}"] = f"='情景参数'!B{row}"
        scenario[f"C{row}"] = f"='基础参数'!B3*(1+B{row})"
        scenario[f"D{row}"] = f"='基础参数'!B2*(1+'基础参数'!B6*B{row})"
        scenario[f"E{row}"] = f"=C{row}*D{row}"
        scenario[f"F{row}"] = f"=D{row}*'基础参数'!B4"
        scenario[f"G{row}"] = f"=E{row}-F{row}"
        scenario[f"H{row}"] = f"=G{row}-'基础参数'!B5"
    finish(scenario)
    percent(scenario, ["B"], 2, 4)
    money(scenario, ["C", "E", "F", "G", "H"], 2, 4)

    matrix = wb.create_sheet("敏感性矩阵")
    matrix["A1"] = "营业利润敏感性：调价幅度 × 额外业务量变动"
    matrix.merge_cells("A1:F1")
    matrix.append(["额外业务量变动", -0.10, -0.05, 0, 0.05, 0.10])
    for row, volume_change in enumerate([-0.10, 0, 0.10], 3):
        matrix[f"A{row}"] = volume_change
        for col in range(2, 7):
            letter = get_column_letter(col)
            matrix[f"{letter}{row}"] = (
                f"=('基础参数'!B2*(1+'基础参数'!B6*{letter}$2)*(1+$A{row}))*"
                f"('基础参数'!B3*(1+{letter}$2)-'基础参数'!B4)-'基础参数'!B5"
            )
            matrix[f"{letter}{row}"].number_format = '¥#,##0.00;[Red]-¥#,##0.00'
    header(matrix, 2)
    body(matrix, 3)
    for cell in matrix[2][1:]:
        cell.number_format = "0%"
    for row in range(3, 6):
        matrix[f"A{row}"].number_format = "0%"
    matrix.freeze_panes = "B3"
    matrix.sheet_view.showGridLines = False
    matrix.conditional_formatting.add("B3:F5", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"))
    for col in range(1, 7):
        matrix.column_dimensions[get_column_letter(col)].width = 18
    save(wb, output_dir / "物流定价情景分析.xlsx")



def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--output-dir", required=True)
    args = parser.parse_args()
    input_dir = Path(args.input_dir).resolve()
    output_dir = Path(args.output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    solve(input_dir, output_dir)


if __name__ == "__main__":
    main()
