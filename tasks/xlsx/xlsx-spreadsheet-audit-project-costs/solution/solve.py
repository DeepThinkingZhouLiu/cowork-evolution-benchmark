
from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "17365D"
BLUE = "4472C4"
GREEN = "70AD47"
ORANGE = "ED7D31"
RED = "C00000"
WHITE = "FFFFFF"
LIGHT_BLUE = "D9EAF7"
THIN = Side(style="thin", color="B7B7B7")


def header(ws, row=1):
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.font = Font(name="微软雅黑", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ws.row_dimensions[row].height = 28


def body(ws, start_row=2):
    for row in ws.iter_rows(min_row=start_row):
        for cell in row:
            cell.font = Font(name="微软雅黑", size=10, color="222222")
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def finish(ws, header_row=1, freeze=None, autofilter=True):
    header(ws, header_row)
    body(ws, header_row + 1)
    ws.freeze_panes = freeze or f"A{header_row + 1}"
    if autofilter:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.sheet_view.showGridLines = False
    for col in range(1, ws.max_column + 1):
        vals = [str(ws.cell(r, col).value or "") for r in range(1, min(ws.max_row, 80) + 1)]
        width = min(max(max((len(v) for v in vals), default=8) * 1.2 + 2, 10), 30)
        ws.column_dimensions[get_column_letter(col)].width = width


def table(ws, name, ref=None):
    ref = ref or f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    obj = Table(displayName=name, ref=ref)
    obj.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(obj)


def money(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = '¥#,##0.00;[Red]-¥#,##0.00'


def percent(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = "0.00%"


def save(wb, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(path)


def solve(input_dir, output_dir):
    ledger_wb = load_workbook(input_dir / "项目成本台账_待审计.xlsx", data_only=True)
    limit_wb = load_workbook(input_dir / "采购合同限额.xlsx", data_only=True)
    headers = [cell.value for cell in ledger_wb.active[1]]
    ledger = [dict(zip(headers, row)) for row in ledger_wb.active.iter_rows(min_row=2, values_only=True)]
    limits = {row[0]: {"category": row[1], "single": row[2], "total": row[3]} for row in limit_wb.active.iter_rows(min_row=2, values_only=True)}
    findings = []
    number = 1

    def add(risk, level, tx, evidence, suggestion, amount):
        nonlocal number
        findings.append([f"F-{number:03d}", risk, level, tx, evidence, suggestion, amount])
        number += 1

    by_invoice = defaultdict(list)
    by_contract = defaultdict(list)
    for item in ledger:
        by_invoice[item["发票号"]].append(item)
        by_contract[item["合同号"]].append(item)
        limit = limits[item["合同号"]]
        if item["金额"] > limit["single"]:
            add("单笔超限", "高", item["交易编号"], f"金额 {item['金额']:.2f} > 单笔限额 {limit['single']:.2f}", "追加授权审批或冲销超限部分", item["金额"])
        if not str(item["审批人"] or "").strip():
            add("缺失审批人", "中", item["交易编号"], "审批人字段为空", "补齐可核验的审批记录", abs(item["金额"]))
        if item["金额"] < 0:
            add("负数金额", "中", item["交易编号"], f"金额为 {item['金额']:.2f}", "核实是否为合法退款并关联原发票", abs(item["金额"]))
    for invoice, items in sorted(by_invoice.items()):
        if len(items) > 1:
            add("重复发票", "高", ",".join(item["交易编号"] for item in items), f"发票号 {invoice} 出现 {len(items)} 次", "暂停支付并核对原始票据", sum(abs(item["金额"]) for item in items))
    for contract, items in sorted(by_contract.items()):
        total = sum(item["金额"] for item in items)
        limit = limits[contract]["total"]
        if total > limit:
            add("合同累计超限", "高", ",".join(item["交易编号"] for item in items), f"合同 {contract} 累计 {total:.2f} > 限额 {limit:.2f}", "停止新增支付并发起合同变更审批", total - limit)

    wb = Workbook()
    original = wb.active
    original.title = "原始台账"
    original.append(headers)
    for item in ledger:
        original.append([item[h] for h in headers])
    finish(original)
    money(original, ["G"])
    table(original, "OriginalLedger")

    finding_ws = wb.create_sheet("审计发现")
    finding_ws.append(["发现编号", "风险类型", "风险级别", "交易编号", "证据", "建议", "涉及金额"])
    for row in findings:
        finding_ws.append(row)
    finish(finding_ws)
    money(finding_ws, ["G"])
    finding_ws.conditional_formatting.add(f"C2:C{finding_ws.max_row}", CellIsRule(operator="equal", formula=['"高"'], fill=PatternFill("solid", fgColor="F4CCCC")))
    table(finding_ws, "AuditFindings")

    summary = wb.create_sheet("审计汇总")
    grouped = defaultdict(lambda: [0, 0.0])
    for finding in findings:
        grouped[finding[1]][0] += 1
        grouped[finding[1]][1] += finding[6]
    summary.append(["风险类型", "发现数", "涉及金额"])
    for risk in sorted(grouped):
        summary.append([risk, grouped[risk][0], grouped[risk][1]])
    summary.append(["合计", len(findings), sum(item[6] for item in findings)])
    finish(summary)
    money(summary, ["C"])
    save(wb, output_dir / "项目成本审计.xlsx")



def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--output-dir", required=True)
    args = parser.parse_args()
    input_dir = Path(args.input_dir).resolve()
    output_dir = Path(args.output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    solve(input_dir, output_dir)


if __name__ == "__main__":
    main()
