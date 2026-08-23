
from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "17365D"
BLUE = "4472C4"
GREEN = "70AD47"
ORANGE = "ED7D31"
RED = "C00000"
WHITE = "FFFFFF"
LIGHT_BLUE = "D9EAF7"
THIN = Side(style="thin", color="B7B7B7")


def header(ws, row=1):
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.font = Font(name="微软雅黑", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ws.row_dimensions[row].height = 28


def body(ws, start_row=2):
    for row in ws.iter_rows(min_row=start_row):
        for cell in row:
            cell.font = Font(name="微软雅黑", size=10, color="222222")
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def finish(ws, header_row=1, freeze=None, autofilter=True):
    header(ws, header_row)
    body(ws, header_row + 1)
    ws.freeze_panes = freeze or f"A{header_row + 1}"
    if autofilter:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.sheet_view.showGridLines = False
    for col in range(1, ws.max_column + 1):
        vals = [str(ws.cell(r, col).value or "") for r in range(1, min(ws.max_row, 80) + 1)]
        width = min(max(max((len(v) for v in vals), default=8) * 1.2 + 2, 10), 30)
        ws.column_dimensions[get_column_letter(col)].width = width


def table(ws, name, ref=None):
    ref = ref or f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    obj = Table(displayName=name, ref=ref)
    obj.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(obj)


def money(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = '¥#,##0.00;[Red]-¥#,##0.00'


def percent(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = "0.00%"


def save(wb, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(path)


def solve(input_dir, output_dir):
    wb = load_workbook(input_dir / "年度预算_公式损坏.xlsx", data_only=False)
    budget = wb["部门预算"]
    summary = wb["汇总"]
    changes = []

    def repair(ws, cell_ref, new_value, reason):
        old = ws[cell_ref].value
        if old != new_value:
            changes.append([ws.title, cell_ref, old, new_value, reason])
            ws[cell_ref] = new_value

    for row in range(2, 7):
        repair(budget, f"F{row}", f"=SUM(B{row}:E{row})", "修复年度合计范围")
        repair(budget, f"H{row}", f"=G{row}-F{row}", "统一差异口径为获批预算-实际")
    for col in "BCDEFG":
        repair(budget, f"{col}7", f"=SUM({col}2:{col}6)", "修复合计行范围")
    repair(budget, "H7", "=G7-F7", "修复合计差异")
    for row in range(2, 7):
        repair(summary, f"B{row}", f"='部门预算'!F{row}", "修复年度实际引用")
        repair(summary, f"C{row}", f"='部门预算'!G{row}", "修复获批预算引用")
        repair(summary, f"D{row}", f"='部门预算'!H{row}", "修复差异引用")
    changes.append(["部门预算", "D1", "表头白色/非粗体", "深蓝底白色粗体", "恢复表头统一样式"])
    header(budget, 1)
    finish(budget)
    finish(summary)
    for ws in [budget, summary]:
        money(ws, [get_column_letter(col) for col in range(2, ws.max_column + 1)])

    log = wb.create_sheet("修复日志")
    log.append(["工作表", "单元格", "原公式/值", "修复后公式/值", "修复原因"])
    for row in changes:
        escaped = list(row)
        for index in [2, 3]:
            if isinstance(escaped[index], str) and escaped[index].startswith("="):
                escaped[index] = "'" + escaped[index]
        log.append(escaped)
    finish(log)
    save(wb, output_dir / "年度预算修复版.xlsx")



def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--output-dir", required=True)
    args = parser.parse_args()
    input_dir = Path(args.input_dir).resolve()
    output_dir = Path(args.output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    solve(input_dir, output_dir)


if __name__ == "__main__":
    main()
