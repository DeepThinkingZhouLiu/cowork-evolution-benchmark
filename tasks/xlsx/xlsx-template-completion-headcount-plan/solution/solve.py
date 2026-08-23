
from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "17365D"
BLUE = "4472C4"
GREEN = "70AD47"
ORANGE = "ED7D31"
RED = "C00000"
WHITE = "FFFFFF"
LIGHT_BLUE = "D9EAF7"
THIN = Side(style="thin", color="B7B7B7")


def header(ws, row=1):
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.font = Font(name="微软雅黑", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ws.row_dimensions[row].height = 28


def body(ws, start_row=2):
    for row in ws.iter_rows(min_row=start_row):
        for cell in row:
            cell.font = Font(name="微软雅黑", size=10, color="222222")
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def finish(ws, header_row=1, freeze=None, autofilter=True):
    header(ws, header_row)
    body(ws, header_row + 1)
    ws.freeze_panes = freeze or f"A{header_row + 1}"
    if autofilter:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.sheet_view.showGridLines = False
    for col in range(1, ws.max_column + 1):
        vals = [str(ws.cell(r, col).value or "") for r in range(1, min(ws.max_row, 80) + 1)]
        width = min(max(max((len(v) for v in vals), default=8) * 1.2 + 2, 10), 30)
        ws.column_dimensions[get_column_letter(col)].width = width


def table(ws, name, ref=None):
    ref = ref or f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    obj = Table(displayName=name, ref=ref)
    obj.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(obj)


def money(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = '¥#,##0.00;[Red]-¥#,##0.00'


def percent(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = "0.00%"


def save(wb, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(path)


def solve(input_dir, output_dir):
    wb = load_workbook(input_dir / "人员编制计划模板.xlsx", data_only=False)
    ws = wb["岗位需求"]
    rows = []
    with (input_dir / "招聘需求.csv").open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    for index, item in enumerate(rows, 4):
        ws[f"A{index}"] = item["部门"]
        ws[f"B{index}"] = item["岗位"]
        ws[f"C{index}"] = int(item["目标编制"])
        ws[f"D{index}"] = int(item["现有人数"])
        ws[f"E{index}"] = f"=MAX(C{index}-D{index},0)"
        ws[f"F{index}"] = datetime.strptime(item["期望到岗日"], "%Y-%m-%d").date()
        ws[f"G{index}"] = f'=IF(OR(E{index}>2,F{index}<DATE(2026,10,1)),"高","常规")'
        ws[f"F{index}"].number_format = "yyyy-mm-dd"
    body(ws, 4)
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:G{3 + len(rows)}"
    ws.conditional_formatting.add(
        f"G4:G{3 + len(rows)}", CellIsRule(operator="equal", formula=['"高"'], fill=PatternFill("solid", fgColor="F4CCCC"))
    )
    summary = wb.create_sheet("部门汇总")
    summary.append(["部门", "目标编制", "现有人数", "招聘缺口"])
    departments = sorted({row["部门"] for row in rows})
    last = 3 + len(rows)
    for index, department in enumerate(departments, 2):
        summary[f"A{index}"] = department
        summary[f"B{index}"] = f'=SUMIF(\'岗位需求\'!A4:A{last},A{index},\'岗位需求\'!C4:C{last})'
        summary[f"C{index}"] = f'=SUMIF(\'岗位需求\'!A4:A{last},A{index},\'岗位需求\'!D4:D{last})'
        summary[f"D{index}"] = f'=SUMIF(\'岗位需求\'!A4:A{last},A{index},\'岗位需求\'!E4:E{last})'
    finish(summary)
    save(wb, output_dir / "人员编制计划完成版.xlsx")



def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--output-dir", required=True)
    args = parser.parse_args()
    input_dir = Path(args.input_dir).resolve()
    output_dir = Path(args.output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    solve(input_dir, output_dir)


if __name__ == "__main__":
    main()
