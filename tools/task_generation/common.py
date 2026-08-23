"""公开样例任务的机械化生成工具。

这里只生成无隐私的合成输入、参考答案和确定性 Judge。正式题仍需人工复核、
视觉 Judge 校准和多模型难度 Rollout，不能把生成成功等同于正式发布。
"""

from __future__ import annotations

import csv
import hashlib
import json
import os
import subprocess
import sys
from pathlib import Path
from typing import Any, Iterable

import fitz
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt as PptPt


PROJECT_PYTHON = Path(os.environ.get("COWORK_EVO_PYTHON", sys.executable))
REFERENCE_URLS = {
    "officeval": "https://github.com/baidu-frontier-research/OmegaUse-OfficeVal",
    "workspace": "https://github.com/OpenDataBox/Workspace-Bench",
    "cowork": "https://github.com/0717376/cowork_bench",
    "officebench": "https://github.com/zlwang-cs/OfficeBench",
    "docx_redline": "https://github.com/dealfluence/docx-benchmark",
    "docx_editing": "https://github.com/nberk/llm-docx-editing",
    "forte": "https://github.com/AGI-Eval-Official/FORTE",
}
SUPPLIER_CAPABILITIES_BY_FAMILY = {
    "material_to_deck": ["C1", "C2", "C4", "C10", "C11"],
    "outline_to_deck": ["C2", "C4", "C10"],
    "data_to_deck": ["C2", "C4", "C5", "C8", "C10"],
    "template_based_generation": ["C3", "C4", "C10"],
    "content_editing": ["C3", "C10"],
    "visual_editing": ["C3", "C4", "C10"],
    "deck_polishing": ["C3", "C4", "C10"],
    "deck_consolidation": ["C2", "C3", "C9", "C10"],
    "audience_adaptation": ["C2", "C3", "C10"],
    "localization": ["C3", "C9", "C10"],
    "presentation_package": ["C6", "C10", "C12"],
    "deck_audit": ["C3", "C10"],
    "material_to_report": ["C1", "C2", "C10", "C11"],
    "outline_to_document": ["C2", "C10"],
    "template_filling": ["C9", "C10"],
    "editing_rewriting": ["C2", "C9", "C10"],
    "redlining_review": ["C9", "C10"],
    "formatting_fidelity": ["C4", "C9", "C10"],
    "merge_compare": ["C2", "C9", "C10"],
    "document_extraction": ["C7", "C9", "C10", "C11"],
    "document_to_table": ["C7", "C10", "C11"],
    "data_cleaning": ["C7", "C10"],
    "table_transformation": ["C7", "C8", "C10"],
    "formula_authoring": ["C8", "C10"],
    "model_building": ["C5", "C8", "C10"],
    "scenario_analysis": ["C5", "C8", "C10"],
    "pivot_analysis": ["C5", "C8", "C10"],
    "charting": ["C5", "C10"],
    "dashboard": ["C5", "C10"],
    "template_completion": ["C8", "C10"],
    "spreadsheet_repair": ["C8", "C10"],
    "spreadsheet_audit": ["C8", "C10"],
    "multi_workbook_integration": ["C7", "C8", "C10"],
    "analysis_to_report": ["C2", "C5", "C8", "C10", "C12"],
    "extract_analyze_present": ["C2", "C5", "C7", "C8", "C10", "C12"],
    "research_to_deliverables": ["C1", "C2", "C10", "C12", "C13"],
    "shared_source_multi_output": ["C2", "C9", "C10", "C12"],
    "update_propagation": ["C3", "C8", "C9", "C10", "C12"],
    "cross_file_consistency": ["C7", "C8", "C10", "C12"],
    "template_driven_package": ["C3", "C8", "C9", "C10", "C12"],
    "review_and_revision": ["C3", "C5", "C8", "C9", "C10", "C12"],
    "workspace_orchestration": ["C2", "C7", "C8", "C9", "C10", "C12"],
}


def write_text(path: Path, content: str, executable: bool = False) -> None:
    """写入 UTF-8 文本，并按需添加可执行位。"""

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content.rstrip() + "\n", encoding="utf-8")
    if executable:
        path.chmod(path.stat().st_mode | 0o111)


def write_json(path: Path, payload: Any) -> None:
    """稳定地写入中文 JSON。"""

    write_text(path, json.dumps(payload, ensure_ascii=False, indent=2))


def sha256_file(path: Path) -> str:
    """计算单个文件的 SHA256。"""

    digest = hashlib.sha256()
    with path.open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def input_hashes(input_dir: Path) -> list[str]:
    """返回输入目录内全部文件的稳定哈希列表。"""

    return [sha256_file(path) for path in sorted(input_dir.rglob("*")) if path.is_file()]


def write_csv(path: Path, rows: Iterable[Iterable[Any]]) -> None:
    """写入带 BOM 的 CSV，便于 Excel 直接识别中文。"""

    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file)
        writer.writerows(rows)


def _set_docx_font(run: Any, size: int = 10, bold: bool | None = None) -> None:
    run.font.name = "Noto Sans CJK SC"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _style_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    normal = document.styles["Normal"]
    normal.font.name = "Noto Sans CJK SC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    normal.font.size = Pt(10.5)
    for style_name, size, color in (
        ("Title", 22, "17365D"),
        ("Heading 1", 15, "17365D"),
        ("Heading 2", 12, "2F5597"),
    ):
        style = document.styles[style_name]
        style.font.name = "Noto Sans CJK SC"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)


def create_docx(
    path: Path,
    *,
    title: str,
    sections: list[dict[str, Any]],
    subtitle: str | None = None,
    footer: str = "CoworkEvoBench 公开合成样例",
    redlines: list[dict[str, str]] | None = None,
) -> None:
    """创建可编辑、带基本专业样式的 DOCX。"""

    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    _style_document(document)
    title_paragraph = document.add_paragraph(style="Title")
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_docx_font(title_paragraph.add_run(title), 22, True)
    if subtitle:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(subtitle)
        _set_docx_font(run, 10)
        run.font.color.rgb = RGBColor(89, 89, 89)

    for section in sections:
        document.add_heading(section["heading"], level=int(section.get("level", 1)))
        for text in section.get("paragraphs", []):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(5)
            _set_docx_font(paragraph.add_run(str(text)), 10)
        for text in section.get("bullets", []):
            paragraph = document.add_paragraph(style="List Bullet")
            _set_docx_font(paragraph.add_run(str(text)), 10)
        table_spec = section.get("table")
        if table_spec:
            headers = table_spec["headers"]
            rows = table_spec["rows"]
            table = document.add_table(rows=1, cols=len(headers))
            table.style = "Table Grid"
            for index, value in enumerate(headers):
                cell = table.rows[0].cells[index]
                cell.text = str(value)
                cell._tc.get_or_add_tcPr().append(
                    OxmlElement("w:shd")
                )
                cell._tc.tcPr[-1].set(qn("w:fill"), "D9EAF7")
                for run in cell.paragraphs[0].runs:
                    _set_docx_font(run, 9, True)
            for row_values in rows:
                cells = table.add_row().cells
                for index, value in enumerate(row_values):
                    cells[index].text = str(value)
                    for run in cells[index].paragraphs[0].runs:
                        _set_docx_font(run, 9)

    for change in redlines or []:
        paragraph = document.add_paragraph()
        label = paragraph.add_run(f"{change.get('label', '修订')}：")
        _set_docx_font(label, 10, True)
        deletion = OxmlElement("w:del")
        deletion.set(qn("w:author"), "CoworkEvoBench Reviewer")
        deletion.set(qn("w:date"), "2026-08-24T00:00:00Z")
        del_run = OxmlElement("w:r")
        del_text = OxmlElement("w:delText")
        del_text.text = change["old"]
        del_run.append(del_text)
        deletion.append(del_run)
        insertion = OxmlElement("w:ins")
        insertion.set(qn("w:author"), "CoworkEvoBench Reviewer")
        insertion.set(qn("w:date"), "2026-08-24T00:00:00Z")
        ins_run = OxmlElement("w:r")
        ins_text = OxmlElement("w:t")
        ins_text.text = change["new"]
        ins_run.append(ins_text)
        insertion.append(ins_run)
        paragraph._p.append(deletion)
        paragraph._p.append(insertion)

    footer_paragraph = document.sections[0].footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_paragraph.add_run(footer)
    _set_docx_font(footer_run, 8)
    footer_run.font.color.rgb = RGBColor(127, 127, 127)
    document.save(path)


def create_xlsx(path: Path, sheets: list[dict[str, Any]]) -> None:
    """创建带样式、公式和可选图表的 XLSX。"""

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="17365D")
    accent_fill = PatternFill("solid", fgColor="D9EAF7")
    for sheet_spec in sheets:
        worksheet = workbook.create_sheet(sheet_spec["name"])
        for row in sheet_spec.get("rows", []):
            worksheet.append(list(row))
        if worksheet.max_row >= 1:
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
        for coordinate, value in sheet_spec.get("cells", {}).items():
            worksheet[coordinate] = value
        for coordinate, formula in sheet_spec.get("formulas", {}).items():
            worksheet[coordinate] = formula
        for coordinate, number_format in sheet_spec.get("number_formats", {}).items():
            worksheet[coordinate].number_format = number_format
        for row_index in range(2, worksheet.max_row + 1):
            if row_index % 2 == 0:
                for cell in worksheet[row_index]:
                    cell.fill = accent_fill
        for column_cells in worksheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            worksheet.column_dimensions[column_cells[0].column_letter].width = min(
                max(max_length + 2, 11), 34
            )
        worksheet.freeze_panes = sheet_spec.get("freeze_panes", "A2")
        if worksheet.max_row > 1 and worksheet.max_column > 1:
            worksheet.auto_filter.ref = worksheet.dimensions
        chart_spec = sheet_spec.get("chart")
        if chart_spec:
            chart = LineChart() if chart_spec.get("type") == "line" else BarChart()
            chart.title = chart_spec.get("title", "指标概览")
            data = Reference(
                worksheet,
                min_col=int(chart_spec.get("data_col", 2)),
                min_row=1,
                max_row=worksheet.max_row,
            )
            categories = Reference(
                worksheet,
                min_col=int(chart_spec.get("category_col", 1)),
                min_row=2,
                max_row=worksheet.max_row,
            )
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(categories)
            chart.height = 7
            chart.width = 13
            worksheet.add_chart(chart, chart_spec.get("anchor", "H2"))
    workbook.save(path)


def _add_textbox(slide: Any, text: str, x: float, y: float, w: float, h: float, size: int) -> Any:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Noto Sans CJK SC"
    paragraph.font.size = PptPt(size)
    paragraph.font.color.rgb = PptRGBColor(33, 43, 54)
    return box


def create_pptx(path: Path, slides: list[dict[str, Any]]) -> None:
    """创建可编辑、层级清楚的 16:9 PPTX。"""

    path.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    for index, slide_spec in enumerate(slides, 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = PptRGBColor(246, 248, 251)
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, presentation.slide_width, Inches(0.16))
        band.fill.solid()
        band.fill.fore_color.rgb = PptRGBColor(23, 54, 93)
        band.line.fill.background()
        title = _add_textbox(slide, slide_spec["title"], 0.7, 0.45, 11.8, 0.65, 25)
        title.text_frame.paragraphs[0].font.bold = True
        title.text_frame.paragraphs[0].font.color.rgb = PptRGBColor(23, 54, 93)
        y = 1.35
        for bullet in slide_spec.get("bullets", []):
            box = _add_textbox(slide, f"• {bullet}", 0.9, y, 11.4, 0.48, 15)
            box.text_frame.word_wrap = True
            y += 0.62
        metrics = slide_spec.get("metrics", [])
        if metrics:
            count = max(len(metrics), 1)
            width = min(3.4, 11.4 / count - 0.2)
            for metric_index, metric in enumerate(metrics):
                x = 0.8 + metric_index * (width + 0.3)
                card = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    Inches(x),
                    Inches(2.0),
                    Inches(width),
                    Inches(1.8),
                )
                card.fill.solid()
                card.fill.fore_color.rgb = PptRGBColor(217, 234, 247)
                card.line.color.rgb = PptRGBColor(91, 155, 213)
                frame = card.text_frame
                frame.clear()
                value_paragraph = frame.paragraphs[0]
                value_paragraph.text = str(metric[1])
                value_paragraph.alignment = PP_ALIGN.CENTER
                value_paragraph.font.size = PptPt(25)
                value_paragraph.font.bold = True
                value_paragraph.font.color.rgb = PptRGBColor(23, 54, 93)
                label_paragraph = frame.add_paragraph()
                label_paragraph.text = str(metric[0])
                label_paragraph.alignment = PP_ALIGN.CENTER
                label_paragraph.font.size = PptPt(11)
        table_spec = slide_spec.get("table")
        if table_spec:
            rows = [table_spec["headers"], *table_spec["rows"]]
            shape = slide.shapes.add_table(
                len(rows),
                len(rows[0]),
                Inches(0.8),
                Inches(1.55),
                Inches(11.7),
                Inches(min(4.8, 0.48 * len(rows))),
            )
            table = shape.table
            for row_index, values in enumerate(rows):
                for col_index, value in enumerate(values):
                    cell = table.cell(row_index, col_index)
                    cell.text = str(value)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = (
                        PptRGBColor(23, 54, 93)
                        if row_index == 0
                        else PptRGBColor(255, 255, 255)
                    )
                    paragraph = cell.text_frame.paragraphs[0]
                    paragraph.font.name = "Noto Sans CJK SC"
                    paragraph.font.size = PptPt(10)
                    paragraph.font.color.rgb = (
                        PptRGBColor(255, 255, 255)
                        if row_index == 0
                        else PptRGBColor(33, 43, 54)
                    )
                    paragraph.font.bold = row_index == 0
        footer = _add_textbox(slide, f"CoworkEvoBench 公开合成样例  |  {index}", 8.5, 7.0, 4.2, 0.25, 8)
        footer.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
        footer.text_frame.paragraphs[0].font.color.rgb = PptRGBColor(110, 118, 129)
    presentation.save(path)


def create_pdf(path: Path, lines: list[str], title: str = "Source Document") -> None:
    """生成只含开放合成内容的可搜索 PDF。"""

    path.parent.mkdir(parents=True, exist_ok=True)
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    page.insert_text((60, 70), title, fontsize=18)
    y = 110
    for line in lines:
        page.insert_text((60, y), line, fontsize=10)
        y += 20
    document.save(path)
    document.close()


SOLVER_TEMPLATE = r'''#!/usr/bin/env python3
"""隐藏评测侧参考解法：根据任务规格生成标准办公交付物。"""

from __future__ import annotations

import argparse
import html
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt as PptPt

SPEC = json.loads(r__SPEC_JSON__)


def font(run, size=10, bold=None):
    run.font.name = "Noto Sans CJK SC"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def build_docx(output, spec):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
    sec.left_margin, sec.right_margin = Cm(2.2), Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = "Noto Sans CJK SC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    normal.font.size = Pt(10.5)
    for name, size, color in (("Title", 22, "17365D"), ("Heading 1", 15, "17365D"), ("Heading 2", 12, "2F5597")):
        style = doc.styles[name]
        style.font.name = "Noto Sans CJK SC"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run(spec["title"]), 22, True)
    if spec.get("subtitle"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(p.add_run(spec["subtitle"]), 10)
    for section in spec["sections"]:
        doc.add_heading(section["heading"], level=section.get("level", 1))
        for text in section.get("paragraphs", []):
            p = doc.add_paragraph()
            font(p.add_run(str(text)), 10)
        for text in section.get("bullets", []):
            p = doc.add_paragraph(style="List Bullet")
            font(p.add_run(str(text)), 10)
        if section.get("table"):
            table_spec = section["table"]
            table = doc.add_table(rows=1, cols=len(table_spec["headers"]))
            table.style = "Table Grid"
            for col, value in enumerate(table_spec["headers"]):
                table.rows[0].cells[col].text = str(value)
                for run in table.rows[0].cells[col].paragraphs[0].runs:
                    font(run, 9, True)
            for values in table_spec["rows"]:
                cells = table.add_row().cells
                for col, value in enumerate(values):
                    cells[col].text = str(value)
                    for run in cells[col].paragraphs[0].runs:
                        font(run, 9)
    for change in spec.get("redlines", []):
        p = doc.add_paragraph()
        font(p.add_run(change.get("label", "修订") + "："), 10, True)
        deletion = OxmlElement("w:del")
        deletion.set(qn("w:author"), "CoworkEvoBench Reviewer")
        deletion.set(qn("w:date"), "2026-08-24T00:00:00Z")
        run = OxmlElement("w:r")
        text = OxmlElement("w:delText")
        text.text = change["old"]
        run.append(text)
        deletion.append(run)
        insertion = OxmlElement("w:ins")
        insertion.set(qn("w:author"), "CoworkEvoBench Reviewer")
        insertion.set(qn("w:date"), "2026-08-24T00:00:00Z")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = change["new"]
        run.append(text)
        insertion.append(run)
        p._p.append(deletion)
        p._p.append(insertion)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run("CoworkEvoBench 公开合成样例"), 8)
    doc.save(output)


def build_xlsx(output, spec):
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_spec in spec["sheets"]:
        ws = wb.create_sheet(sheet_spec["name"])
        for row in sheet_spec.get("rows", []):
            ws.append(row)
        for coordinate, value in sheet_spec.get("cells", {}).items():
            ws[coordinate] = value
        for coordinate, value in sheet_spec.get("formulas", {}).items():
            ws[coordinate] = value
        if ws.max_row:
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="17365D")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center")
        for column in ws.columns:
            width = max(len(str(cell.value or "")) for cell in column) + 2
            ws.column_dimensions[column[0].column_letter].width = min(max(width, 11), 34)
        ws.freeze_panes = "A2"
        if ws.max_row > 1 and ws.max_column > 1:
            ws.auto_filter.ref = ws.dimensions
        chart_spec = sheet_spec.get("chart")
        if chart_spec:
            chart = LineChart() if chart_spec.get("type") == "line" else BarChart()
            chart.title = chart_spec.get("title", "指标概览")
            data = Reference(ws, min_col=chart_spec.get("data_col", 2), min_row=1, max_row=ws.max_row)
            categories = Reference(ws, min_col=chart_spec.get("category_col", 1), min_row=2, max_row=ws.max_row)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(categories)
            ws.add_chart(chart, chart_spec.get("anchor", "H2"))
    wb.save(output)


def textbox(slide, text, x, y, w, h, size):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.name = "Noto Sans CJK SC"
    p.font.size = PptPt(size)
    p.font.color.rgb = PptRGBColor(33, 43, 54)
    return box


def build_pptx(output, spec):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for index, slide_spec in enumerate(spec["slides"], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = PptRGBColor(246, 248, 251)
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.16))
        band.fill.solid()
        band.fill.fore_color.rgb = PptRGBColor(23, 54, 93)
        band.line.fill.background()
        title = textbox(slide, slide_spec["title"], 0.7, 0.45, 11.8, 0.65, 25)
        title.text_frame.paragraphs[0].font.bold = True
        title.text_frame.paragraphs[0].font.color.rgb = PptRGBColor(23, 54, 93)
        y = 1.35
        for bullet in slide_spec.get("bullets", []):
            textbox(slide, "• " + str(bullet), 0.9, y, 11.4, 0.48, 15)
            y += 0.62
        metrics = slide_spec.get("metrics", [])
        for metric_index, metric in enumerate(metrics):
            x = 0.8 + metric_index * 3.8
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2), Inches(3.4), Inches(1.8))
            card.fill.solid()
            card.fill.fore_color.rgb = PptRGBColor(217, 234, 247)
            card.line.color.rgb = PptRGBColor(91, 155, 213)
            frame = card.text_frame
            frame.clear()
            p = frame.paragraphs[0]
            p.text = str(metric[1])
            p.alignment = PP_ALIGN.CENTER
            p.font.size = PptPt(25)
            p.font.bold = True
            p.font.color.rgb = PptRGBColor(23, 54, 93)
            p = frame.add_paragraph()
            p.text = str(metric[0])
            p.alignment = PP_ALIGN.CENTER
            p.font.size = PptPt(11)
        if slide_spec.get("table"):
            table_spec = slide_spec["table"]
            rows = [table_spec["headers"], *table_spec["rows"]]
            shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(0.8), Inches(1.55), Inches(11.7), Inches(min(4.8, 0.48 * len(rows))))
            for row_index, values in enumerate(rows):
                for col_index, value in enumerate(values):
                    cell = shape.table.cell(row_index, col_index)
                    cell.text = str(value)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = PptRGBColor(23, 54, 93) if row_index == 0 else PptRGBColor(255, 255, 255)
                    p = cell.text_frame.paragraphs[0]
                    p.font.size = PptPt(10)
                    p.font.color.rgb = PptRGBColor(255, 255, 255) if row_index == 0 else PptRGBColor(33, 43, 54)
                    p.font.bold = row_index == 0
        footer = textbox(slide, f"CoworkEvoBench 公开合成样例  |  {index}", 8.5, 7.0, 4.2, 0.25, 8)
        footer.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    prs.save(output)


def build_html(output, spec):
    title = html.escape(str(spec["title"]))
    subtitle = html.escape(str(spec.get("subtitle", "")))
    decision = html.escape(str(spec["decision"]))
    guardrail = html.escape(str(spec["guardrail"]))
    review_date = html.escape(str(spec["review_date"]))
    metrics = "".join(
        '<article class="metric"><strong>{}</strong><span>{}</span></article>'.format(
            html.escape(str(value)), html.escape(str(label))
        )
        for label, value in spec.get("metrics", [])
    )
    evidence = "".join(
        f"<li>{html.escape(str(item))}</li>" for item in spec.get("evidence", [])
    )
    actions = "".join(
        "<tr><td>{}</td><td>{}</td><td>{}</td></tr>".format(
            *(html.escape(str(value)) for value in row)
        )
        for row in spec.get("actions", [])
    )
    document = f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{title}</title>
  <style>
    :root {{ color-scheme: light; font-family: "Noto Sans CJK SC", sans-serif; }}
    body {{ margin: 0; background: #f4f7fb; color: #1f2937; }}
    main {{ max-width: 1040px; margin: 0 auto; padding: 40px 24px 64px; }}
    header {{ color: white; background: #17365d; border-radius: 18px; padding: 28px; }}
    h1 {{ margin: 0 0 8px; font-size: 30px; }}
    .metrics {{ display: grid; grid-template-columns: repeat(3, 1fr); gap: 16px; margin: 22px 0; }}
    .metric, section {{ background: white; border-radius: 14px; padding: 20px; box-shadow: 0 4px 16px #17365d12; }}
    .metric strong {{ display: block; color: #17365d; font-size: 28px; }}
    .metric span {{ color: #52606d; }}
    section {{ margin-top: 18px; }}
    .decision {{ border-left: 6px solid #16a085; }}
    .guardrail {{ border-left: 6px solid #e67e22; }}
    table {{ width: 100%; border-collapse: collapse; }}
    th, td {{ text-align: left; padding: 10px; border-bottom: 1px solid #dbe3ec; }}
    button {{ border: 0; border-radius: 9px; padding: 10px 14px; background: #2f5597; color: white; cursor: pointer; }}
    [hidden] {{ display: none; }}
    @media (max-width: 720px) {{ .metrics {{ grid-template-columns: 1fr; }} }}
  </style>
</head>
<body>
<main>
  <header><h1>{title}</h1><p>{subtitle}</p></header>
  <div class="metrics">{metrics}</div>
  <section class="decision"><h2>决策建议</h2><p>{decision}</p></section>
  <section class="guardrail"><h2>强制护栏</h2><p>{guardrail}</p><p>下一次复审：{review_date}</p></section>
  <section>
    <h2>研究证据</h2>
    <button id="evidence-button" type="button" aria-expanded="false" onclick="toggleEvidence()">展开研究证据</button>
    <ul id="evidence-list" hidden>{evidence}</ul>
  </section>
  <section><h2>扩展行动</h2><table><thead><tr><th>行动</th><th>责任</th><th>时间</th></tr></thead><tbody>{actions}</tbody></table></section>
</main>
<script>
function toggleEvidence() {{
  const list = document.getElementById('evidence-list');
  const button = document.getElementById('evidence-button');
  list.hidden = !list.hidden;
  button.setAttribute('aria-expanded', String(!list.hidden));
  button.textContent = list.hidden ? '展开研究证据' : '收起研究证据';
}}
</script>
</body>
</html>
"""
    output.write_text(document, encoding="utf-8")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    missing = [name for name in SPEC["required_inputs"] if not (args.input_dir / name).exists()]
    if missing:
        raise SystemExit(f"缺少输入文件：{missing}")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for spec in SPEC.get("documents", []):
        build_docx(args.output_dir / spec["filename"], spec)
    for spec in SPEC.get("workbooks", []):
        build_xlsx(args.output_dir / spec["filename"], spec)
    for spec in SPEC.get("presentations", []):
        build_pptx(args.output_dir / spec["filename"], spec)
    for spec in SPEC.get("html_pages", []):
        build_html(args.output_dir / spec["filename"], spec)


if __name__ == "__main__":
    main()
'''


JUDGE_TEMPLATE = r'''#!/usr/bin/env python3
"""任务级确定性 Judge；语义与视觉质量将在正式版叠加独立 Judge。"""

from __future__ import annotations

import argparse
import json
import math
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


def ratio(checks):
    return 1.0 if not checks else sum(bool(item) for item in checks) / len(checks)


def inspect_docx(path):
    doc = Document(path)
    text = "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml")
    root = ElementTree.fromstring(xml)
    xml_text = "".join(root.itertext())
    return {
        "text": text + "\n" + xml_text,
        "tables": len(doc.tables),
        "headings": sum(p.style and p.style.name.startswith("Heading") for p in doc.paragraphs),
        "styled": sum(bool(p.style and p.style.name != "Normal") for p in doc.paragraphs),
        "redlines": xml.count(b"<w:ins") + xml.count(b"<w:del"),
    }


def inspect_xlsx(path):
    wb = load_workbook(path, data_only=False)
    text = []
    charts = 0
    cells = {}
    for ws in wb.worksheets:
        charts += len(ws._charts)
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    text.append(str(cell.value))
                    cells[f"{ws.title}!{cell.coordinate}"] = cell.value
    result = {"text": "\n".join(text), "sheets": wb.sheetnames, "charts": charts, "cells": cells}
    wb.close()
    return result


def inspect_pptx(path):
    prs = Presentation(path)
    text = []
    shape_count = 0
    for slide in prs.slides:
        shape_count += len(slide.shapes)
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                text.append(shape.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    text.extend(cell.text for cell in row.cells)
    return {"text": "\n".join(text), "slides": len(prs.slides), "shapes": shape_count}


def inspect_html(path):
    text = path.read_text(encoding="utf-8")
    lowered = text.lower()
    if "<html" not in lowered or "</html>" not in lowered:
        raise ValueError("HTML 根元素不完整")
    return {
        "text": text,
        "buttons": lowered.count("<button"),
        "scripts": lowered.count("<script"),
    }


def equal_value(actual, expected):
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        try:
            return math.isclose(float(actual), float(expected), rel_tol=1e-6, abs_tol=1e-6)
        except (TypeError, ValueError):
            return False
    return str(actual) == str(expected)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("artifact_dir", type=Path)
    parser.add_argument("--expectations", type=Path, default=Path(__file__).with_name("expectations.json"))
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    expected = json.loads(args.expectations.read_text(encoding="utf-8"))
    inspected, open_checks = {}, []
    for filename in expected["required_files"]:
        path = args.artifact_dir / filename
        if not path.is_file():
            open_checks.append(False)
            continue
        try:
            suffix = path.suffix.lower()
            if suffix == ".docx":
                inspected[filename] = inspect_docx(path)
            elif suffix == ".xlsx":
                inspected[filename] = inspect_xlsx(path)
            elif suffix == ".pptx":
                inspected[filename] = inspect_pptx(path)
            elif suffix == ".html":
                inspected[filename] = inspect_html(path)
            else:
                inspected[filename] = {"text": path.read_text(encoding="utf-8")}
            open_checks.append(True)
        except Exception:
            open_checks.append(False)

    content_checks, exact_checks, structure_checks, presentation_checks = [], [], [], []
    preservation_checks, penalty_checks = [], []
    for filename, file_expected in expected.get("files", {}).items():
        info = inspected.get(filename, {"text": ""})
        text = info.get("text", "")
        content_checks.extend(term in text for term in file_expected.get("contains", []))
        exact_checks.extend(term in text for term in file_expected.get("exact_terms", []))
        preservation_checks.extend(term in text for term in file_expected.get("preserve_terms", []))
        penalty_checks.extend(term not in text for term in file_expected.get("excludes", []))
        if "min_tables" in file_expected:
            structure_checks.append(info.get("tables", 0) >= file_expected["min_tables"])
        if "min_headings" in file_expected:
            structure_checks.append(info.get("headings", 0) >= file_expected["min_headings"])
        if "min_redlines" in file_expected:
            structure_checks.append(info.get("redlines", 0) >= file_expected["min_redlines"])
        if "sheet_names" in file_expected:
            structure_checks.append(all(name in info.get("sheets", []) for name in file_expected["sheet_names"]))
        if "min_slides" in file_expected:
            structure_checks.append(info.get("slides", 0) >= file_expected["min_slides"])
        if "max_slides" in file_expected:
            structure_checks.append(info.get("slides", 0) <= file_expected["max_slides"])
        if "min_charts" in file_expected:
            presentation_checks.append(info.get("charts", 0) >= file_expected["min_charts"])
        if "min_shapes" in file_expected:
            presentation_checks.append(info.get("shapes", 0) >= file_expected["min_shapes"])
        if "min_styled" in file_expected:
            presentation_checks.append(info.get("styled", 0) >= file_expected["min_styled"])
        if "min_buttons" in file_expected:
            presentation_checks.append(info.get("buttons", 0) >= file_expected["min_buttons"])
        if "min_scripts" in file_expected:
            presentation_checks.append(info.get("scripts", 0) >= file_expected["min_scripts"])
        for coordinate, value in file_expected.get("cells", {}).items():
            exact_checks.append(equal_value(info.get("cells", {}).get(coordinate), value))
        for coordinate, value in file_expected.get("formulas", {}).items():
            exact_checks.append(info.get("cells", {}).get(coordinate) == value)

    consistency_checks = []
    for rule in expected.get("shared_terms", []):
        consistency_checks.append(all(rule["term"] in inspected.get(name, {}).get("text", "") for name in rule["files"]))
    consistency_checks.extend(preservation_checks)

    scores = {
        "R001": ratio(open_checks),
        "R002": ratio(content_checks),
        "R003": ratio(exact_checks),
        "R004": ratio(structure_checks),
        "R005": ratio(presentation_checks),
        "R006": ratio(consistency_checks),
        "R007": ratio(penalty_checks),
    }
    positive_weights = {"R001": 2, "R002": 4, "R003": 4, "R004": 2, "R005": 2, "R006": 2}
    reward = sum(scores[key] * weight for key, weight in positive_weights.items()) / sum(positive_weights.values())
    reward = max(0.0, reward - (1.0 - scores["R007"]) * 3 / sum(positive_weights.values()))
    if scores["R001"] < 1:
        reward = 0.0
    passed = reward >= expected.get("pass_threshold", 0.7)
    result = {
        "task_id": expected["task_id"],
        "judge_version": "cowork-evo-deterministic-v0.1",
        "reward": round(reward, 6),
        "passed": passed,
        "criterion_results": [
            {"criterion_id": key, "score": round(value, 6), "method": "deterministic"}
            for key, value in scores.items()
        ],
    }
    output = json.dumps(result, ensure_ascii=False, indent=2)
    print(output)
    if args.output:
        args.output.write_text(output + "\n", encoding="utf-8")
    raise SystemExit(0 if passed else 1)


if __name__ == "__main__":
    main()
'''


def _toml_list(values: Iterable[str]) -> str:
    return "[" + ", ".join(json.dumps(value, ensure_ascii=False) for value in values) + "]"


def write_task_package(
    task_root: Path,
    *,
    task_id: str,
    title: str,
    description: str,
    instruction: str,
    metadata: dict[str, Any],
    solve_spec: dict[str, Any],
    expectations: dict[str, Any],
    source_urls: list[str],
) -> None:
    """补齐一个自包含 Harbor 任务包并生成 Golden。"""

    for relative in (
        "environment",
        "data/input_files",
        "solution",
        "tests",
        "metadata",
    ):
        (task_root / relative).mkdir(parents=True, exist_ok=True)

    full_task_id = f"cowork-evo/{task_id}"
    task_metadata = {
        "task_id": full_task_id,
        "domain": metadata["domain"],
        "task_family": metadata["task_family"],
        "taxonomy_version": "v0.1",
        "operation_types": metadata["operation_types"],
        "input_modalities": metadata["input_modalities"],
        "dependency_scope": metadata["dependency_scope"],
        "constraint_profiles": metadata["constraint_profiles"],
        "deliverable_types": metadata["deliverable_types"],
        "output_filenames": expectations["required_files"],
        "capability_ids": metadata["capability_ids"],
        "supplier_capability_ids": SUPPLIER_CAPABILITIES_BY_FAMILY[
            metadata["task_family"]
        ],
        "sector": metadata["sector"],
        "occupation": metadata["occupation"],
        "difficulty_tier": metadata["difficulty_tier"],
        "multimodal": metadata.get("multimodal", False),
        "web_search_required": False,
        "language": "zh-CN",
        "judge_types": ["deterministic", *metadata.get("additional_judges", [])],
        "estimated_minutes": metadata["estimated_minutes"],
        "provenance_id": task_id,
        "leakage_group_id": f"public-synthetic-{metadata['task_family']}-001",
        "publishability": "public",
        "notes": "公开合成端到端样例；正式发布前仍需人工与视觉 Judge 复核。",
    }
    write_json(task_root / "metadata/task.json", task_metadata)
    write_text(task_root / "instruction.md", instruction)
    write_text(
        task_root / "task.toml",
        f'''schema_version = "1.4"
artifacts = ["/app/output", "/logs/artifacts/output"]

[task]
name = "{full_task_id}"
version = "0.1.0"
description = {json.dumps(description, ensure_ascii=False)}
keywords = {_toml_list([metadata["domain"], metadata["task_family"], "office", "public-sample"])}

[metadata]
domain = "{metadata["domain"]}"
task_family = "{metadata["task_family"]}"
taxonomy_version = "v0.1"
operation_types = {_toml_list(metadata["operation_types"])}
input_modalities = {_toml_list(metadata["input_modalities"])}
dependency_scope = "{metadata["dependency_scope"]}"
constraint_profiles = {_toml_list(metadata["constraint_profiles"])}
deliverable_types = {_toml_list(metadata["deliverable_types"])}
output_filenames = {_toml_list(expectations["required_files"])}
capability_ids = {_toml_list(metadata["capability_ids"])}
supplier_capability_ids = {_toml_list(SUPPLIER_CAPABILITIES_BY_FAMILY[metadata["task_family"]])}
sector = {json.dumps(metadata["sector"], ensure_ascii=False)}
occupation = {json.dumps(metadata["occupation"], ensure_ascii=False)}
difficulty_tier = "{metadata["difficulty_tier"]}"
multimodal = {str(metadata.get("multimodal", False)).lower()}
web_search_required = false
language = "zh-CN"
judge_types = {_toml_list(["deterministic", *metadata.get("additional_judges", [])])}
estimated_minutes = {metadata["estimated_minutes"]}
provenance_id = "{task_id}"
leakage_group_id = "public-synthetic-{metadata["task_family"]}-001"
publishability = "public"

[agent]
timeout_sec = 1800.0

[environment]
build_timeout_sec = 1800.0
cpus = 2
memory_mb = 4096
storage_mb = 12288
gpus = 0
allow_internet = false
mcp_servers = []
workdir = "/app"

[verifier]
timeout_sec = 1800.0
''',
    )
    rubric = {
        "task_id": full_task_id,
        "rubric_version": "cowork-evo-rubric-v0.1",
        "pass_threshold": 0.7,
        "criteria": [
            {
                "id": "R001",
                "description": f"{title}要求的全部交付文件存在、格式正确且可以打开。",
                "weight": 2,
                "type": "hurdle",
                "dimension": "file_integrity",
                "method": "deterministic",
                "method_params": {},
                "score_cap_on_failure": 0,
            },
            {
                "id": "R002",
                "description": "交付物完整覆盖用户指定的核心事实、结论与行动项。",
                "weight": 4,
                "type": "positive",
                "dimension": "instruction_following",
                "method": "hybrid",
                "method_params": {},
            },
            {
                "id": "R003",
                "description": "关键数字、表格单元格、公式或版本事实与输入材料一致。",
                "weight": 4,
                "type": "positive",
                "dimension": "data_accuracy",
                "method": "deterministic",
                "method_params": {},
            },
            {
                "id": "R004",
                "description": "章节、页面、工作表和表格结构满足任务要求且逻辑完整。",
                "weight": 2,
                "type": "positive",
                "dimension": "structure",
                "method": "deterministic",
                "method_params": {},
            },
            {
                "id": "R005",
                "description": "交付物具有清楚的视觉层级、基本格式和可直接使用的呈现质量。",
                "weight": 2,
                "type": "positive",
                "dimension": "presentation",
                "method": "hybrid",
                "method_params": {},
            },
            {
                "id": "R006",
                "description": "要求保留的事实没有丢失，多文件中的共享数字与结论保持一致。",
                "weight": 2,
                "type": "positive",
                "dimension": "consistency",
                "method": "deterministic",
                "method_params": {},
            },
            {
                "id": "R007",
                "description": "不得残留模板占位符、过期事实或输入无法支持的关键陈述。",
                "weight": -3,
                "type": "penalty",
                "dimension": "factuality",
                "method": "hybrid",
                "method_params": {},
            },
        ],
    }
    write_json(task_root / "rubric.json", rubric)

    solve_json = json.dumps(solve_spec, ensure_ascii=False)
    solve_source = SOLVER_TEMPLATE.replace("r__SPEC_JSON__", "r'''" + solve_json + "'''")
    write_text(task_root / "solution/solve.py", solve_source, executable=True)
    write_text(
        task_root / "solution/solve.sh",
        '''#!/usr/bin/env bash
set -euo pipefail
task_root="$(cd "$(dirname "${BASH_SOURCE[0]}")/.." && pwd)"
output_dir="${1:-${OUTPUT_DIR:-/app/output}}"
mkdir -p "$output_dir"
python3 "$task_root/solution/solve.py" --input-dir "$task_root/data/input_files" --output-dir "$output_dir"
''',
        executable=True,
    )
    write_text(task_root / "tests/judge.py", JUDGE_TEMPLATE, executable=True)
    expectations = {"task_id": full_task_id, "pass_threshold": 0.7, **expectations}
    write_json(task_root / "tests/expectations.json", expectations)
    write_text(
        task_root / "tests/test.sh",
        '''#!/usr/bin/env bash
set -euo pipefail
task_root="$(cd "$(dirname "${BASH_SOURCE[0]}")/.." && pwd)"
artifact_dir="${1:-${ARTIFACT_DIR:-$task_root/solution}}"
python3 "$task_root/tests/judge.py" "$artifact_dir"
''',
        executable=True,
    )
    write_text(
        task_root / "environment/Dockerfile",
        '''FROM python:3.11-slim
RUN pip install --no-cache-dir openpyxl==3.1.5 python-pptx==1.0.2 python-docx==1.2.0 lxml==6.1.2 Pillow==11.3.0
WORKDIR /app
''',
    )
    provenance = {
        "provenance_id": task_id,
        "source_type": "synthetic",
        "source_uri": source_urls[0] if source_urls else None,
        "acquired_at": "2026-08-24T00:00:00+08:00",
        "license": "CC0-1.0（本项目合成输入）；参考代码许可见 003-Reference 索引",
        "privacy_status": "通过",
        "human_rewritten": True,
        "synthetic": True,
        "publishability": "public",
        "source_sha256": input_hashes(task_root / "data/input_files"),
        "notes": "只参考任务形态与 Judge 设计，未复制上游用户内容。参考：" + "；".join(source_urls),
    }
    write_json(task_root / "metadata/provenance.json", provenance)

    result = subprocess.run(
        [
            str(PROJECT_PYTHON),
            str(task_root / "solution/solve.py"),
            "--input-dir",
            str(task_root / "data/input_files"),
            "--output-dir",
            str(task_root / "solution"),
        ],
        check=False,
        text=True,
        capture_output=True,
    )
    if result.returncode:
        raise RuntimeError(f"{task_id} Golden 生成失败：{result.stderr}")
