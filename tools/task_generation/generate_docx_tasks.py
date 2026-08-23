"""生成 8 个 Word 任务家族的公开端到端样例。"""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.shared import Pt

from common import (
    REFERENCE_URLS,
    create_docx,
    create_pdf,
    write_csv,
    write_json,
    write_task_package,
    write_text,
)


ROOT = Path(__file__).resolve().parents[2]
TASKS_ROOT = ROOT / "tasks/docx"


def metadata(
    family: str,
    *,
    operations: list[str],
    modalities: list[str],
    scope: str,
    constraints: list[str],
    capabilities: list[str],
    sector: str,
    occupation: str,
    difficulty: str = "intermediate",
    minutes: int = 45,
    outputs: list[str] | None = None,
) -> dict:
    return {
        "domain": "docx",
        "task_family": family,
        "operation_types": operations,
        "input_modalities": modalities,
        "dependency_scope": scope,
        "constraint_profiles": constraints,
        "deliverable_types": outputs or ["docx"],
        "capability_ids": capabilities,
        "sector": sector,
        "occupation": occupation,
        "difficulty_tier": difficulty,
        "estimated_minutes": minutes,
        "multimodal": "image" in modalities,
        "additional_judges": ["semantic", "visual"],
    }


def instruction(title: str, context: str, inputs: list[str], outputs: list[str], requirements: list[str]) -> str:
    lines = [
        f"# {title}",
        "",
        context,
        "",
        "## 输入文件",
        "",
        *[f"- `{name}`" for name in inputs],
        "",
        "## 交付要求",
        "",
        *[f"- {item}" for item in requirements],
        "",
        "## 输出位置",
        "",
        *[f"- `/app/output/{name}`" for name in outputs],
        "",
        "不得联网，不得把来源材料中没有出现的数字写成事实。交付文件必须可编辑并可由 Word 或 LibreOffice 正常打开。",
    ]
    return "\n".join(lines)


def build_material_to_report() -> None:
    task_id = "docx-001-material-to-report"
    root = TASKS_ROOT / task_id
    input_dir = root / "data/input_files"
    write_csv(
        input_dir / "regional_sales.csv",
        [
            ["区域", "Q2收入_元", "同比增长", "重点客户续约率"],
            ["华东", 1260000, "18%", "94%"],
            ["华南", 980000, "9%", "88%"],
            ["华北", 760000, "-4%", "82%"],
            ["合计", 3000000, "8.7%", "88%"],
        ],
    )
    write_text(
        input_dir / "customer_interviews.md",
        """# 客户访谈摘录

- 华东客户认可交付速度，但希望增加月度经营复盘。
- 华南客户认为培训材料分散，建议建立统一知识库。
- 华北两家客户反馈上线支持响应慢，是续约风险的主要原因。
""",
    )
    write_json(
        input_dir / "report_requirements.json",
        {"audience": "经营管理会", "period": "2026 Q2", "owner": "商业运营部", "max_pages": 6},
    )
    output = "Q2区域经营复盘.docx"
    solve_spec = {
        "required_inputs": ["regional_sales.csv", "customer_interviews.md", "report_requirements.json"],
        "documents": [
            {
                "filename": output,
                "title": "2026 Q2 区域经营复盘",
                "subtitle": "经营管理会材料｜商业运营部",
                "sections": [
                    {
                        "heading": "执行摘要",
                        "paragraphs": [
                            "Q2 总收入为 3,000,000 元，同比增长 8.7%。华东以 1,260,000 元和 18% 的同比增速领先；华北同比下降 4%，应列为优先修复区域。"
                        ],
                    },
                    {
                        "heading": "核心指标",
                        "table": {
                            "headers": ["区域", "Q2收入（元）", "同比增长", "续约率"],
                            "rows": [
                                ["华东", "1,260,000", "18%", "94%"],
                                ["华南", "980,000", "9%", "88%"],
                                ["华北", "760,000", "-4%", "82%"],
                                ["合计", "3,000,000", "8.7%", "88%"],
                            ],
                        },
                    },
                    {
                        "heading": "客户反馈与诊断",
                        "bullets": [
                            "华东：增加月度经营复盘，巩固高续约率。",
                            "华南：整合培训材料，建立统一知识库。",
                            "华北：上线支持响应慢，是当前最主要的续约风险。",
                        ],
                    },
                    {
                        "heading": "下一步行动",
                        "table": {
                            "headers": ["优先级", "行动", "负责人", "完成时间"],
                            "rows": [
                                ["P0", "建立华北上线支持值班机制", "客户成功负责人", "2026-07-15"],
                                ["P1", "上线统一培训知识库", "交付运营", "2026-07-31"],
                                ["P1", "启动华东月度经营复盘", "区域经理", "2026-07-10"],
                            ],
                        },
                    },
                ],
            }
        ],
    }
    expectations = {
        "required_files": [output],
        "files": {
            output: {
                "contains": ["执行摘要", "核心指标", "客户反馈与诊断", "下一步行动", "华北"],
                "exact_terms": ["3,000,000", "8.7%", "1,260,000", "18%", "-4%", "82%"],
                "preserve_terms": ["月度经营复盘", "统一知识库", "上线支持响应慢"],
                "excludes": ["{{", "待补充", "同比增长 20%"],
                "min_tables": 2,
                "min_headings": 4,
                "min_styled": 5,
            }
        },
    }
    write_task_package(
        root,
        task_id=task_id,
        title="区域经营复盘报告",
        description="从区域销售表和访谈摘录生成管理层经营复盘。",
        instruction=instruction(
            "区域经营复盘报告",
            "你是商业运营经理，需要把 Q2 区域销售数据和客户访谈整理成可直接用于经营管理会的 Word 报告。",
            ["regional_sales.csv", "customer_interviews.md", "report_requirements.json"],
            [output],
            [
                "报告包含执行摘要、核心指标、客户反馈与诊断、下一步行动四个章节。",
                "准确呈现三大区域及合计数字，明确指出领先区域和风险区域。",
                "行动表至少包含优先级、行动、负责人和完成时间。",
                "使用清楚的标题层级和专业表格，不超过 6 页。",
            ],
        ),
        metadata=metadata(
            "material_to_report",
            operations=["create", "analyze"],
            modalities=["plain_text", "spreadsheet", "multi_file"],
            scope="multi_file",
            constraints=["content_fidelity", "numeric_accuracy", "traceability", "professional_readiness"],
            capabilities=["K1", "K2", "K5", "K6", "K10"],
            sector="企业服务",
            occupation="商业运营经理",
            difficulty="advanced",
            minutes=60,
        ),
        solve_spec=solve_spec,
        expectations=expectations,
        source_urls=[REFERENCE_URLS["officeval"], REFERENCE_URLS["forte"]],
    )


def build_outline_to_document() -> None:
    task_id = "docx-002-outline-to-document"
    root = TASKS_ROOT / task_id
    input_dir = root / "data/input_files"
    write_text(
        input_dir / "onboarding_outline.md",
        """# 新员工入职手册大纲

1. 第一天准备
2. 信息安全
3. 费用报销
4. 30/60/90 天目标
5. 常见问题
""",
    )
    write_json(
        input_dir / "company_policies.json",
        {
            "work_hours": "09:30-18:30",
            "security_training_deadline": "入职后 5 个工作日内",
            "expense_deadline": "费用发生后 30 天内",
            "probation_review_day": 90,
            "helpdesk": "helpdesk@example.test",
        },
    )
    output = "新员工入职手册.docx"
    solve_spec = {
        "required_inputs": ["onboarding_outline.md", "company_policies.json"],
        "documents": [
            {
                "filename": output,
                "title": "新员工入职手册",
                "subtitle": "People Operations｜2026 版",
                "sections": [
                    {"heading": "第一天准备", "bullets": ["工作时间：09:30-18:30", "完成账号激活、设备检查和直属经理见面。"]},
                    {"heading": "信息安全", "paragraphs": ["请在入职后 5 个工作日内完成信息安全培训，不得共享账号或将公司材料上传到未授权服务。"]},
                    {"heading": "费用报销", "paragraphs": ["所有费用须在发生后 30 天内提交，并附合规票据与业务说明。"]},
                    {
                        "heading": "30/60/90 天目标",
                        "table": {
                            "headers": ["节点", "目标", "检查方式"],
                            "rows": [["30 天", "理解团队目标和关键流程", "经理一对一"], ["60 天", "独立承担核心工作", "阶段复盘"], ["90 天", "完成试用期评审", "正式评审"]],
                        },
                    },
                    {"heading": "常见问题", "bullets": ["IT 支持：helpdesk@example.test", "制度疑问：联系 People Operations。"]},
                ],
            }
        ],
    }
    expectations = {
        "required_files": [output],
        "files": {
            output: {
                "contains": ["第一天准备", "信息安全", "费用报销", "30/60/90 天目标", "常见问题"],
                "exact_terms": ["09:30-18:30", "5 个工作日内", "30 天内", "90 天", "helpdesk@example.test"],
                "excludes": ["{{", "TBD", "7 个工作日内"],
                "min_tables": 1,
                "min_headings": 5,
                "min_styled": 6,
            }
        },
    }
    write_task_package(
        root,
        task_id=task_id,
        title="新员工入职手册",
        description="把简要大纲和制度事实扩写成完整入职手册。",
        instruction=instruction(
            "把大纲扩写成新员工入职手册",
            "你是 People Operations 专员。公司只有一份章节大纲和结构化制度事实，需要形成一份新人拿到即可使用的手册。",
            ["onboarding_outline.md", "company_policies.json"],
            [output],
            ["保留大纲的五个一级章节并补成完整内容。", "所有时间、期限和联系方式严格以 JSON 为准。", "30/60/90 天目标使用表格表达。", "语气友好、明确，不编造福利或制度。"],
        ),
        metadata=metadata(
            "outline_to_document",
            operations=["create"],
            modalities=["plain_text", "multi_file"],
            scope="multi_file",
            constraints=["content_fidelity", "style_consistency", "professional_readiness"],
            capabilities=["K1", "K2", "K6", "K9"],
            sector="人力资源",
            occupation="People Operations 专员",
            minutes=45,
        ),
        solve_spec=solve_spec,
        expectations=expectations,
        source_urls=[REFERENCE_URLS["officeval"], REFERENCE_URLS["docx_editing"]],
    )


def build_template_filling() -> None:
    task_id = "docx-003-template-filling"
    root = TASKS_ROOT / task_id
    input_dir = root / "data/input_files"
    create_docx(
        input_dir / "供应商评审模板.docx",
        title="供应商评审结论",
        sections=[
            {"heading": "基本信息", "paragraphs": ["供应商：{{vendor_name}}", "服务范围：{{service_scope}}", "合同金额：{{contract_value}}"]},
            {"heading": "服务水平", "paragraphs": ["请根据 SLA 明细表填写关键指标。"]},
            {"heading": "评审结论", "paragraphs": ["结论：{{decision}}", "生效日期：{{effective_date}}"]},
        ],
    )
    write_json(
        input_dir / "vendor_facts.json",
        {
            "vendor_name": "明川云科技有限公司",
            "service_scope": "客户数据备份与灾备演练",
            "contract_value": "480,000 元/年",
            "decision": "有条件通过",
            "effective_date": "2026-09-01",
            "condition": "首月完成一次全量恢复演练",
        },
    )
    write_csv(input_dir / "sla.csv", [["指标", "承诺值"], ["月度可用性", "99.95%"], ["P1 响应", "15 分钟"], ["恢复时间目标", "4 小时"]])
    output = "供应商评审结论_明川云.docx"
    solve_spec = {
        "required_inputs": ["供应商评审模板.docx", "vendor_facts.json", "sla.csv"],
        "documents": [
            {
                "filename": output,
                "title": "供应商评审结论",
                "sections": [
                    {"heading": "基本信息", "table": {"headers": ["字段", "内容"], "rows": [["供应商", "明川云科技有限公司"], ["服务范围", "客户数据备份与灾备演练"], ["合同金额", "480,000 元/年"]]}},
                    {"heading": "服务水平", "table": {"headers": ["指标", "承诺值"], "rows": [["月度可用性", "99.95%"], ["P1 响应", "15 分钟"], ["恢复时间目标", "4 小时"]]}},
                    {"heading": "评审结论", "paragraphs": ["结论：有条件通过。生效日期：2026-09-01。", "前置条件：首月完成一次全量恢复演练。"]},
                ],
            }
        ],
    }
    expectations = {
        "required_files": [output],
        "files": {
            output: {
                "contains": ["基本信息", "服务水平", "评审结论", "有条件通过", "全量恢复演练"],
                "exact_terms": ["明川云科技有限公司", "480,000 元/年", "99.95%", "15 分钟", "4 小时", "2026-09-01"],
                "excludes": ["{{", "}}", "待填写"],
                "min_tables": 2,
                "min_headings": 3,
                "min_styled": 4,
            }
        },
    }
    write_task_package(
        root,
        task_id=task_id,
        title="供应商评审模板填充",
        description="将供应商事实和 SLA 准确填入公司 Word 模板。",
        instruction=instruction(
            "完成供应商评审结论",
            "你是采购风控专员，需要使用既有模板出具供应商评审结论。模板占位符必须全部替换，不能改变章节顺序。",
            ["供应商评审模板.docx", "vendor_facts.json", "sla.csv"],
            [output],
            ["填入供应商、服务范围、合同金额、结论和生效日期。", "把三项 SLA 做成清楚表格。", "明确写出有条件通过的前置条件。", "不得残留任何双花括号占位符。"],
        ),
        metadata=metadata(
            "template_filling",
            operations=["edit", "transform"],
            modalities=["document", "spreadsheet", "plain_text", "multi_file"],
            scope="multi_file",
            constraints=["content_fidelity", "layout_fidelity", "style_consistency", "preservation"],
            capabilities=["K1", "K7", "K9", "K10"],
            sector="采购与供应链",
            occupation="采购风控专员",
            difficulty="advanced",
            minutes=55,
        ),
        solve_spec=solve_spec,
        expectations=expectations,
        source_urls=[REFERENCE_URLS["docx_editing"], REFERENCE_URLS["officeval"]],
    )


def build_editing_rewriting() -> None:
    task_id = "docx-004-editing-rewriting"
    root = TASKS_ROOT / task_id
    input_dir = root / "data/input_files"
    create_docx(
        input_dir / "客户沟通初稿.docx",
        title="续约沟通",
        sections=[
            {"heading": "正文", "paragraphs": ["嗨，最近系统其实还行吧。我们想让你们赶紧续约，价格还是 320,000 元，服务期到 2027-06-30。上个月发生了 2 次 P1，我们已经搞定了。下周三前给回复哈。"]}
        ],
    )
    write_text(
        input_dir / "品牌写作规范.md",
        """# 品牌写作规范

- 语气：专业、坦诚、以客户价值为中心。
- 必须承认问题并写出改进动作，不推诿。
- 金额、日期、事件次数必须保留原值。
- 结尾给出清楚但不施压的下一步。
""",
    )
    output = "客户续约沟通稿.docx"
    solve_spec = {
        "required_inputs": ["客户沟通初稿.docx", "品牌写作规范.md"],
        "documents": [
            {
                "filename": output,
                "title": "客户续约沟通稿",
                "sections": [
                    {"heading": "主题", "paragraphs": ["关于下一服务周期续约与稳定性改进计划"]},
                    {"heading": "沟通正文", "paragraphs": ["感谢贵团队在本服务周期中的合作。下一周期合同金额为 320,000 元，服务期至 2027-06-30。", "我们重视上月发生的 2 次 P1 事件。当前已完成根因复盘，并新增值班升级和变更前检查两项措施。", "如方便，希望在下周三前确认是否安排 30 分钟沟通；我们会根据贵团队意见完善续约方案。"]},
                    {"heading": "内部发送检查", "bullets": ["金额、日期和事件次数已复核。", "语气专业、坦诚，未弱化稳定性问题。"]},
                ],
            }
        ],
    }
    expectations = {
        "required_files": [output],
        "files": {
            output: {
                "contains": ["客户续约沟通稿", "根因复盘", "值班升级", "变更前检查", "30 分钟沟通"],
                "exact_terms": ["320,000 元", "2027-06-30", "2 次 P1", "下周三前"],
                "preserve_terms": ["续约", "稳定性"],
                "excludes": ["其实还行吧", "赶紧续约", "搞定了", "给回复哈"],
                "min_headings": 3,
                "min_styled": 4,
            }
        },
    }
    write_task_package(
        root,
        task_id=task_id,
        title="客户沟通专业改写",
        description="按品牌规范重写不专业的续约沟通稿并保留事实。",
        instruction=instruction(
            "改写客户续约沟通稿",
            "你是客户成功负责人。现有初稿语气随意，还弱化了稳定性问题；请按品牌写作规范改成可直接发送给客户的正式版本。",
            ["客户沟通初稿.docx", "品牌写作规范.md"],
            [output],
            ["完整保留金额、服务截止日期、P1 次数和回复时间。", "坦诚承认问题，并写出根因复盘和值班升级、变更前检查等动作。", "删除口语化、施压式表达。", "使用主题、正文和内部发送检查三个清楚部分。"],
        ),
        metadata=metadata(
            "editing_rewriting",
            operations=["edit"],
            modalities=["document", "plain_text", "multi_file"],
            scope="multi_file",
            constraints=["content_fidelity", "preservation", "style_consistency", "professional_readiness"],
            capabilities=["K2", "K6", "K7", "K9"],
            sector="企业服务",
            occupation="客户成功负责人",
            minutes=35,
        ),
        solve_spec=solve_spec,
        expectations=expectations,
        source_urls=[REFERENCE_URLS["officeval"], REFERENCE_URLS["docx_editing"]],
    )


def build_redlining_review() -> None:
    task_id = "docx-005-redlining-review"
    root = TASKS_ROOT / task_id
    input_dir = root / "data/input_files"
    create_docx(
        input_dir / "云服务协议_原稿.docx",
        title="云服务协议",
        sections=[
            {"heading": "第 4 条 服务水平", "paragraphs": ["服务可用性不低于 99.5%，P1 事件响应时间为 60 分钟。"]},
            {"heading": "第 7 条 数据保留", "paragraphs": ["合同终止后，客户数据保留 30 天。"]},
            {"heading": "第 9 条 责任限制", "paragraphs": ["年度累计赔偿责任不超过最近一个月服务费。"]},
        ],
    )
    write_json(
        input_dir / "法务修改意见.json",
        {
            "availability": {"old": "99.5%", "new": "99.9%"},
            "p1_response": {"old": "60 分钟", "new": "30 分钟"},
            "retention": {"old": "30 天", "new": "90 天"},
            "liability": {"old": "最近一个月服务费", "new": "最近三个月服务费"},
            "instruction": "保留修订痕迹，不接受静默覆盖。",
        },
    )
    output = "云服务协议_修订稿.docx"
    redlines = [
        {"label": "第 4 条可用性", "old": "99.5%", "new": "99.9%"},
        {"label": "第 4 条 P1 响应", "old": "60 分钟", "new": "30 分钟"},
        {"label": "第 7 条数据保留", "old": "30 天", "new": "90 天"},
        {"label": "第 9 条责任上限", "old": "最近一个月服务费", "new": "最近三个月服务费"},
    ]
    solve_spec = {
        "required_inputs": ["云服务协议_原稿.docx", "法务修改意见.json"],
        "documents": [
            {
                "filename": output,
                "title": "云服务协议（修订稿）",
                "subtitle": "修订日期：2026-08-24",
                "sections": [
                    {"heading": "审阅说明", "paragraphs": ["本稿按法务修改意见形成，所有实质变更均保留插入与删除标记。"]},
                    {"heading": "未修改条款", "paragraphs": ["除下列四项外，其余条款保持原稿不变。"]},
                ],
                "redlines": redlines,
            }
        ],
    }
    expectations = {
        "required_files": [output],
        "files": {
            output: {
                "contains": ["云服务协议（修订稿）", "审阅说明", "所有实质变更均保留插入与删除标记"],
                "exact_terms": ["99.5%", "99.9%", "60 分钟", "30 分钟", "30 天", "90 天", "最近一个月服务费", "最近三个月服务费"],
                "preserve_terms": ["其余条款保持原稿不变"],
                "excludes": ["{{", "接受全部风险"],
                "min_redlines": 8,
                "min_headings": 2,
                "min_styled": 3,
            }
        },
    }
    write_task_package(
        root,
        task_id=task_id,
        title="云服务协议红线审阅",
        description="按照法务意见修改协议并保留 OOXML 修订痕迹。",
        instruction=instruction(
            "按法务意见修订云服务协议",
            "你是法务运营专员。请把结构化修改意见应用到协议原稿，必须保留可追踪的删除与插入痕迹，不能只生成一份静默覆盖后的干净版本。",
            ["云服务协议_原稿.docx", "法务修改意见.json"],
            [output],
            ["准确完成可用性、P1 响应、数据保留和责任上限四项修改。", "每项同时保留原值删除标记和新值插入标记。", "明确说明其余条款不变。", "文件必须仍是可编辑、可打开的 DOCX。"],
        ),
        metadata=metadata(
            "redlining_review",
            operations=["edit", "audit"],
            modalities=["document", "plain_text", "multi_file"],
            scope="multi_file",
            constraints=["content_fidelity", "preservation", "traceability", "professional_readiness"],
            capabilities=["K2", "K7", "K9", "K12"],
            sector="法律服务",
            occupation="法务运营专员",
            difficulty="expert",
            minutes=70,
        ),
        solve_spec=solve_spec,
        expectations=expectations,
        source_urls=[REFERENCE_URLS["docx_redline"], REFERENCE_URLS["docx_editing"]],
    )


def build_formatting_fidelity() -> None:
    task_id = "docx-006-formatting-fidelity"
    root = TASKS_ROOT / task_id
    input_dir = root / "data/input_files"
    input_dir.mkdir(parents=True, exist_ok=True)
    draft = Document()
    p = draft.add_paragraph("北辰零售 7 月经营简报")
    p.runs[0].font.size = Pt(9)
    draft.add_paragraph("核心结论")
    draft.add_paragraph("7 月销售额 5,480,000 元，较 6 月增长 6.2%。会员销售贡献 61%。")
    draft.add_paragraph("门店表现")
    draft.add_paragraph("浦东店 1,920,000 元；滨江店 1,760,000 元；天府店 1,800,000 元。")
    draft.add_paragraph("下月动作")
    draft.add_paragraph("优化会员唤醒；控制缺货率；复盘高毛利品类陈列。")
    draft.save(input_dir / "经营简报_待排版.docx")
    write_json(
        input_dir / "brand_guide.json",
        {"primary_color": "#17365D", "accent_color": "#5B9BD5", "title": "22pt", "heading_1": "15pt", "body": "10.5pt", "footer": "北辰零售｜内部使用"},
    )
    output = "经营简报_品牌版.docx"
    solve_spec = {
        "required_inputs": ["经营简报_待排版.docx", "brand_guide.json"],
        "documents": [
            {
                "filename": output,
                "title": "北辰零售 7 月经营简报",
                "subtitle": "内部经营材料｜2026 年 7 月",
                "sections": [
                    {"heading": "核心结论", "paragraphs": ["7 月销售额 5,480,000 元，较 6 月增长 6.2%。会员销售贡献 61%。"]},
                    {"heading": "门店表现", "table": {"headers": ["门店", "销售额（元）"], "rows": [["浦东店", "1,920,000"], ["滨江店", "1,760,000"], ["天府店", "1,800,000"]]}},
                    {"heading": "下月动作", "bullets": ["优化会员唤醒", "控制缺货率", "复盘高毛利品类陈列"]},
                ],
            }
        ],
    }
    expectations = {
        "required_files": [output],
        "files": {
            output: {
                "contains": ["核心结论", "门店表现", "下月动作", "优化会员唤醒", "控制缺货率"],
                "exact_terms": ["5,480,000 元", "6.2%", "61%", "1,920,000", "1,760,000", "1,800,000"],
                "preserve_terms": ["北辰零售 7 月经营简报", "复盘高毛利品类陈列"],
                "excludes": ["{{", "待排版"],
                "min_tables": 1,
                "min_headings": 3,
                "min_styled": 4,
            }
        },
    }
    write_task_package(
        root,
        task_id=task_id,
        title="经营简报格式保真",
        description="在不改变原文事实的前提下应用品牌格式和信息层级。",
        instruction=instruction(
            "把经营简报排成品牌版",
            "你是品牌运营专员。原文内容正确但完全没有层级，请按品牌指南重新排版，同时确保任何数字和行动项都不丢失。",
            ["经营简报_待排版.docx", "brand_guide.json"],
            [output],
            ["建立标题、一级标题、正文和页脚层级。", "门店表现改成两列表格，三个门店数字保持不变。", "下月动作使用项目符号。", "不得改写销售额、增长率、会员贡献和行动项。"],
        ),
        metadata=metadata(
            "formatting_fidelity",
            operations=["edit"],
            modalities=["document", "plain_text", "multi_file"],
            scope="multi_file",
            constraints=["layout_fidelity", "style_consistency", "preservation", "numeric_accuracy"],
            capabilities=["K7", "K8", "K9", "K12"],
            sector="零售",
            occupation="品牌运营专员",
            minutes=40,
        ),
        solve_spec=solve_spec,
        expectations=expectations,
        source_urls=[REFERENCE_URLS["officeval"], REFERENCE_URLS["docx_editing"]],
    )


def build_merge_compare() -> None:
    task_id = "docx-007-merge-compare"
    root = TASKS_ROOT / task_id
    input_dir = root / "data/input_files"
    create_docx(
        input_dir / "差旅政策_v1.docx",
        title="差旅政策 v1",
        sections=[
            {"heading": "住宿", "paragraphs": ["一线城市住宿上限 600 元/晚，其他城市 400 元/晚。"]},
            {"heading": "交通", "paragraphs": ["高铁行程不超过 5 小时应选择二等座。"]},
            {"heading": "审批", "paragraphs": ["出差前由直属经理审批。"]},
        ],
    )
    create_docx(
        input_dir / "差旅政策_v2.docx",
        title="差旅政策 v2",
        sections=[
            {"heading": "住宿", "paragraphs": ["一线城市住宿上限 700 元/晚，其他城市 450 元/晚。"]},
            {"heading": "交通", "paragraphs": ["高铁行程不超过 6 小时应选择二等座。"]},
            {"heading": "审批", "paragraphs": ["出差前由直属经理审批；预算超过 5,000 元需部门负责人复核。"]},
            {"heading": "碳排披露", "paragraphs": ["跨省出差须在报销时填写交通方式。"]},
        ],
    )
    outputs = ["差旅政策_合并版.docx", "差旅政策_版本差异.docx"]
    solve_spec = {
        "required_inputs": ["差旅政策_v1.docx", "差旅政策_v2.docx"],
        "documents": [
            {
                "filename": outputs[0],
                "title": "差旅政策（合并版）",
                "subtitle": "生效日期：2026-09-01",
                "sections": [
                    {"heading": "住宿", "paragraphs": ["一线城市住宿上限 700 元/晚，其他城市 450 元/晚。"]},
                    {"heading": "交通", "paragraphs": ["高铁行程不超过 6 小时应选择二等座。"]},
                    {"heading": "审批", "paragraphs": ["出差前由直属经理审批；预算超过 5,000 元需部门负责人复核。"]},
                    {"heading": "碳排披露", "paragraphs": ["跨省出差须在报销时填写交通方式。"]},
                ],
            },
            {
                "filename": outputs[1],
                "title": "差旅政策版本差异",
                "sections": [
                    {
                        "heading": "v1 与 v2 差异矩阵",
                        "table": {
                            "headers": ["条款", "v1", "v2/合并决定"],
                            "rows": [
                                ["一线城市住宿", "600 元/晚", "700 元/晚"],
                                ["其他城市住宿", "400 元/晚", "450 元/晚"],
                                ["高铁时长", "5 小时", "6 小时"],
                                ["大额预算审批", "无", "超过 5,000 元需复核"],
                                ["碳排披露", "无", "新增交通方式披露"],
                            ],
                        },
                    },
                    {"heading": "合并原则", "paragraphs": ["合并版以 v2 为最新政策事实，并通过差异表保留 v1 的历史信息。"]},
                ],
            },
        ],
    }
    expectations = {
        "required_files": outputs,
        "files": {
            outputs[0]: {
                "contains": ["住宿", "交通", "审批", "碳排披露", "部门负责人复核"],
                "exact_terms": ["700 元/晚", "450 元/晚", "6 小时", "5,000 元"],
                "excludes": ["600 元/晚", "400 元/晚", "不超过 5 小时"],
                "min_headings": 4,
                "min_styled": 5,
            },
            outputs[1]: {
                "contains": ["v1 与 v2 差异矩阵", "600 元/晚", "400 元/晚", "5 小时", "新增交通方式披露"],
                "exact_terms": ["700 元/晚", "450 元/晚", "6 小时", "5,000 元"],
                "excludes": ["{{"],
                "min_tables": 1,
                "min_headings": 2,
                "min_styled": 3,
            },
        },
        "shared_terms": [{"term": "700 元/晚", "files": outputs}, {"term": "5,000 元", "files": outputs}],
    }
    write_task_package(
        root,
        task_id=task_id,
        title="差旅政策版本合并",
        description="比较两个 Word 版本并交付合并政策与差异说明。",
        instruction=instruction(
            "合并两个差旅政策版本",
            "你是公司治理专员。v2 是最新政策，但管理层还需要明确知道它相对 v1 改了什么。请同时交付一份干净合并版和一份差异说明。",
            ["差旅政策_v1.docx", "差旅政策_v2.docx"],
            outputs,
            ["合并版以 v2 的最新值为准，章节完整。", "差异说明用表格逐项列出住宿、交通、审批和新增披露变化。", "两个文件中的最新政策数字必须一致。", "差异说明保留 v1 的旧值，合并版不得混入旧值。"],
        ),
        metadata=metadata(
            "merge_compare",
            operations=["audit", "edit", "package"],
            modalities=["document", "multi_file"],
            scope="multi_file",
            constraints=["content_fidelity", "preservation", "traceability", "cross_artifact_consistency"],
            capabilities=["K2", "K7", "K10", "K11", "K12"],
            sector="公司治理",
            occupation="公司治理专员",
            difficulty="advanced",
            minutes=60,
        ),
        solve_spec=solve_spec,
        expectations=expectations,
        source_urls=[REFERENCE_URLS["docx_redline"], REFERENCE_URLS["workspace"]],
    )


def build_document_extraction() -> None:
    task_id = "docx-008-document-extraction"
    root = TASKS_ROOT / task_id
    input_dir = root / "data/input_files"
    create_pdf(
        input_dir / "invoice.pdf",
        [
            "Invoice No: MC-2026-0718",
            "Vendor: Mingchuan Cloud Technology Co., Ltd.",
            "Service: Disaster recovery drill - Q3",
            "Subtotal: CNY 96,000.00",
            "Tax: CNY 5,760.00",
            "Total: CNY 101,760.00",
            "Payment due: 2026-09-15",
        ],
        title="SERVICE INVOICE",
    )
    write_text(
        input_dir / "acceptance_record.txt",
        """验收单号：AC-2026-0831
验收日期：2026-08-31
验收结论：通过
范围：Q3 灾备恢复演练
遗留事项：在 2026-09-10 前补交演练录像索引
负责人：周敏
""",
    )
    write_csv(input_dir / "purchase_order.csv", [["采购单号", "供应商", "含税金额", "预算科目"], ["PO-2026-0442", "明川云科技有限公司", 101760, "IT-BCP"]])
    output = "付款审核摘要.docx"
    solve_spec = {
        "required_inputs": ["invoice.pdf", "acceptance_record.txt", "purchase_order.csv"],
        "documents": [
            {
                "filename": output,
                "title": "付款审核摘要",
                "subtitle": "采购单 PO-2026-0442",
                "sections": [
                    {"heading": "单据核对", "table": {"headers": ["项目", "发票/采购单/验收", "核对结果"], "rows": [["供应商", "明川云科技有限公司 / Mingchuan Cloud Technology Co., Ltd.", "一致"], ["服务范围", "Q3 灾备恢复演练", "一致"], ["含税金额", "101,760.00 元", "一致"], ["验收结论", "通过", "满足付款条件"]]}},
                    {"heading": "关键编号与日期", "bullets": ["发票：MC-2026-0718", "采购单：PO-2026-0442", "验收单：AC-2026-0831", "付款到期日：2026-09-15"]},
                    {"heading": "付款建议", "paragraphs": ["建议按 101,760.00 元办理付款。验收已通过，但应由负责人周敏在 2026-09-10 前补交演练录像索引并归档。"]},
                    {"heading": "风险提示", "paragraphs": ["补交录像索引属于归档遗留项，不影响当前验收结论；付款后仍需跟踪关闭。"]},
                ],
            }
        ],
    }
    expectations = {
        "required_files": [output],
        "files": {
            output: {
                "contains": ["单据核对", "关键编号与日期", "付款建议", "风险提示", "建议按 101,760.00 元办理付款"],
                "exact_terms": ["MC-2026-0718", "PO-2026-0442", "AC-2026-0831", "101,760.00", "2026-09-15", "2026-09-10", "周敏"],
                "preserve_terms": ["演练录像索引", "验收已通过"],
                "excludes": ["无需补交", "拒绝付款", "{{"],
                "min_tables": 1,
                "min_headings": 4,
                "min_styled": 5,
            }
        },
    }
    write_task_package(
        root,
        task_id=task_id,
        title="跨单据付款审核摘要",
        description="从 PDF 发票、验收记录和采购单抽取事实并形成付款审核文档。",
        instruction=instruction(
            "生成付款审核摘要",
            "你是财务共享中心审核员。请联合发票 PDF、验收记录和采购单核对供应商、服务、金额、编号和日期，并形成一份可归档的付款审核摘要。",
            ["invoice.pdf", "acceptance_record.txt", "purchase_order.csv"],
            [output],
            ["用表格呈现单据核对结果。", "准确写出三个单据编号、含税总额和付款到期日。", "区分验收已通过与仍需补交录像索引两个事实。", "给出明确付款建议和后续跟踪责任。"],
        ),
        metadata=metadata(
            "document_extraction",
            operations=["transform", "analyze", "create"],
            modalities=["document", "spreadsheet", "plain_text", "multi_file"],
            scope="multi_file",
            constraints=["content_fidelity", "numeric_accuracy", "traceability", "professional_readiness"],
            capabilities=["K1", "K2", "K5", "K6", "K10"],
            sector="财务与采购",
            occupation="财务共享中心审核员",
            difficulty="advanced",
            minutes=55,
        ),
        solve_spec=solve_spec,
        expectations=expectations,
        source_urls=[REFERENCE_URLS["officebench"], REFERENCE_URLS["cowork"]],
    )


def main() -> None:
    builders = [
        build_material_to_report,
        build_outline_to_document,
        build_template_filling,
        build_editing_rewriting,
        build_redlining_review,
        build_formatting_fidelity,
        build_merge_compare,
        build_document_extraction,
    ]
    for builder in builders:
        builder()
        print(f"已生成：{builder.__name__}")
    print(f"共生成 {len(builders)} 个 Word 公开端到端样例。")


if __name__ == "__main__":
    main()
