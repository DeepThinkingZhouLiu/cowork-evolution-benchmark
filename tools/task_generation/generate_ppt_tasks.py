#!/usr/bin/env python3
"""生成 CoworkEvoBench 的 12 个 PPT 端到端任务。

这些任务中的业务材料均为专门撰写的合成数据，不包含任何真实个人或企业信息。
参考仓库只用于学习任务拆分、原子 Rubric 和原始文件/修改后文件的配对方式；
本脚本不读取或执行任何参考仓库代码。
"""

from __future__ import annotations

import csv
import hashlib
import json
import pprint
import stat
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt as PptPt

from common import SUPPLIER_CAPABILITIES_BY_FAMILY


REPO_ROOT = Path(__file__).resolve().parents[2]
TASKS_ROOT = REPO_ROOT / "tasks" / "ppt"
ACQUIRED_AT = "2026-08-24T12:00:00+08:00"

REFERENCE_URLS = {
    "PPTArena": "https://github.com/michaelofengenden/PPTArena",
    "PPT-Eval": "https://github.com/microsoft/ppteval",
    "PPTC": "https://github.com/gydpku/PPTC",
    "OmegaUse-OfficeVal": "https://github.com/baidu-frontier-research/OmegaUse-OfficeVal",
    "Workspace-Bench": "https://github.com/OpenDataBox/Workspace-Bench",
}

OPERATION_TYPES_BY_FAMILY = {
    "material_to_deck": ["create", "analyze"],
    "outline_to_deck": ["create"],
    "data_to_deck": ["create", "analyze"],
    "template_based_generation": ["create"],
    "content_editing": ["edit"],
    "visual_editing": ["edit"],
    "deck_polishing": ["edit"],
    "deck_consolidation": ["transform"],
    "audience_adaptation": ["edit", "transform"],
    "localization": ["edit", "transform"],
    "presentation_package": ["package"],
    "deck_audit": ["audit"],
}

INPUT_MODALITY_BY_KIND = {
    "csv": "spreadsheet",
    "docx": "document",
    "jpg": "image",
    "json": "plain_text",
    "md": "plain_text",
    "pdf": "document",
    "png": "image",
    "pptx": "presentation",
    "txt": "plain_text",
    "xlsx": "spreadsheet",
}


def _slide(
    kind: str,
    title: str,
    *,
    subtitle: str = "",
    bullets: list[str] | None = None,
    kpis: list[list[str]] | None = None,
    chart: dict[str, Any] | None = None,
    table: list[list[str]] | None = None,
    timeline: list[list[str]] | None = None,
    callout: str = "",
    note: str = "",
) -> dict[str, Any]:
    return {
        "kind": kind,
        "title": title,
        "subtitle": subtitle,
        "bullets": bullets or [],
        "kpis": kpis or [],
        "chart": chart,
        "table": table,
        "timeline": timeline or [],
        "callout": callout,
        "note": note,
    }


TASKS: list[dict[str, Any]] = [
    {
        "id": "ppt-001-community-canteen-retrospective",
        "family": "material_to_deck",
        "title": "社区食堂试点复盘汇报",
        "sector": "公共服务",
        "occupation": "项目运营经理",
        "difficulty": "advanced",
        "estimated": 90,
        "capabilities": ["K1", "K2", "K4", "K7", "K9"],
        "accent": "E45B35",
        "author": "枫桥街道社区服务中心（合成）",
        "inputs": [
            {
                "name": "社区食堂试点记录.docx",
                "kind": "docx",
                "title": "长者友好社区食堂试点记录",
                "paragraphs": [
                    "试点周期：2026年4月1日至6月30日，覆盖枫桥、河滨两个社区。",
                    "三个月累计服务 8,460 人次，满意度 4.7/5，长者复购率 92%。",
                    "用餐结构：午餐占 64%，晚餐占 26%，其他时段占 10%。",
                    "三类高频建议：增加低盐套餐、周三增设快取窗、对行动不便者试行配送。",
                    "项目组建议在 90 天内扩容，但需街道确认预算、场地与配送试点边界。",
                ],
                "table": [
                    ["反馈主题", "样本数", "建议动作"],
                    ["低盐套餐", "126", "每日不少于2款"],
                    ["周三排队", "89", "11:30–12:15增开快取窗"],
                    ["助餐配送", "57", "先50户小范围试行"],
                ],
            },
            {
                "name": "月度客流与成本.xlsx",
                "kind": "xlsx",
                "sheets": [
                    {
                        "name": "月度统计",
                        "rows": [
                            ["月份", "服务人次", "单均成本(元)", "满意度"],
                            ["4月", 2480, 15.8, 4.6],
                            ["5月", 2810, 15.2, 4.7],
                            ["6月", 3170, 14.9, 4.8],
                            ["合计/均值", 8460, 15.3, 4.7],
                        ],
                    }
                ],
            },
            {"name": "社区食堂场景.png", "kind": "png", "palette": ["F9E0D1", "E45B35", "264653"]},
        ],
        "source_markers": ["8,460", "4.7/5", "92%", "64%", "90 天"],
        "deliverables": [{"name": "社区食堂试点复盘.pptx", "type": "pptx"}],
        "slides": [
            _slide("cover", "长者友好社区食堂｜试点复盘", subtitle="2026年4–6月 · 枫桥/河滨社区", callout="决策版"),
            _slide("kpi", "三个月试点：覆盖与满意度双达标", kpis=[["8,460", "服务人次"], ["4.7/5", "平均满意度"], ["92%", "长者复购率"]], bullets=["服务量连续两月增长", "单均成本由15.8元降至14.9元"]),
            _slide("chart", "午餐贡献 64%，6月客流达 3,170 人次", chart={"categories": ["4月", "5月", "6月"], "series": [["服务人次", [2480, 2810, 3170]]]}, bullets=["周三 11:30–12:15 为最明显峰值", "优先通过快取窗分流，不急于扩建后厨"]),
            _slide("table", "三类反馈直接对应下一阶段动作", table=[["需求", "证据", "下一步"], ["低盐套餐", "126条", "每日2款"], ["减少排队", "89条", "周三快取窗"], ["助餐配送", "57条", "先50户试行"]], callout="先做可逆的小步实验"),
            _slide("timeline", "90 天扩容路线图", timeline=[["0–30天", "低盐菜单 + 快取窗"], ["31–60天", "50户配送试行"], ["61–90天", "根据复购/成本决定扩容"]], callout="每30天复盘一次"),
            _slide("closing", "需街道确认的三项决策", bullets=["批准90天试行预算上限18万元", "协调社区前厅作为周三快取区", "同意配送仅覆盖50户且不收集敏感健康信息"], callout="建议：有条件通过"),
        ],
        "content_tokens": ["长者友好社区食堂", "低盐套餐", "快取窗", "50户", "有条件通过"],
        "numeric_tokens": ["8,460", "4.7/5", "92%", "64%", "3,170", "90 天"],
        "preserve_tokens": ["枫桥/河滨社区", "决策版"],
        "forbidden_tokens": ["真实姓名", "身份证号", "患病名单"],
        "object_mins": {"charts": 1, "tables": 1, "pictures": 1, "shapes": 45},
        "references": ["PPT-Eval/README.md", "OmegaUse-OfficeVal/data/rubrics-en/officeval_005.json"],
        "requirements": ["将分散的记录与表格数据综合为6页决策汇报", "数字必须与输入一致，不得编造长者个人信息", "用图表表达客流趋势，并明确90天决策点"],
    },
    {
        "id": "ppt-002-nebula-launch-outline",
        "family": "outline_to_deck",
        "title": "星澜 S1 新品发布会",
        "sector": "工业科技",
        "occupation": "产品市场经理",
        "difficulty": "intermediate",
        "estimated": 75,
        "capabilities": ["K1", "K3", "K7", "K9"],
        "accent": "6C5CE7",
        "author": "星澜科技（合成）",
        "inputs": [
            {"name": "星澜S1发布会大纲.md", "kind": "text", "content": "# 星澜 S1 发布会\n\n受众：工厂设备部与信息化负责人。\n\n- 行业问题：传统巡检平均耗时45分钟。\n- 核心价值：零接触巡检，单线部署12分钟。\n- 可用性：99.95%。\n- 试点结果：误报率从8.2%降到2.6%。\n- 上市日：2026年10月15日。\n- 路线：内测、伙伴试用、正式发布三阶段。\n"},
            {"name": "品牌配色.json", "kind": "json", "content": {"brand": "星澜科技", "primary": "#181B3A", "accent": "#6C5CE7", "secondary": "#00CEC9", "font": "Noto Sans CJK SC"}},
            {"name": "星澜S1产品示意.png", "kind": "png", "palette": ["181B3A", "6C5CE7", "00CEC9"]},
        ],
        "source_markers": ["星澜 S1", "零接触巡检", "12分钟", "99.95%", "2026年10月15日"],
        "deliverables": [{"name": "星澜S1新品发布会.pptx", "type": "pptx"}],
        "slides": [
            _slide("cover", "星澜 S1｜让巡检从人找问题变成问题找人", subtitle="工业设备智能巡检终端 · 2026秋季发布", callout="2026.10.15"),
            _slide("bullets", "45 分钟的人工巡检，是一笔被忽略的停机风险", bullets=["纸面记录难以追溯，异常通常延迟上报", "高危点位依赖人工靠近，检查频率受限", "多厂区口径不一，管理者无法及时比较"], callout="今天的核心问题：如何在不停线的前提下更早发现异常？"),
            _slide("kpi", "零接触巡检：12 分钟完成单线部署", kpis=[["12分钟", "单线部署"], ["99.95%", "平台可用性"], ["2.6%", "试点误报率"]], bullets=["无需改造现有PLC与传感器", "本地分析 + 异常摘要上传"]),
            _slide("timeline", "从采集到闭环，异常在一个工作台完成", timeline=[["01 感知", "非接触采集"], ["02 判断", "边缘模型识别"], ["03 协同", "工单与责任人"], ["04 复盘", "可追溯的趋势"]]),
            _slide("chart", "试点误报率从 8.2% 降至 2.6%", chart={"categories": ["上线前", "第2周", "第6周"], "series": [["误报率", [8.2, 4.1, 2.6]]]}, bullets=["数据来自同一条包装线的连续6周记录", "不将试点结果外推为全行业保证"]),
            _slide("table", "三阶段发布，每一步都有可回退门槛", table=[["阶段", "时间", "成功门槛"], ["内测", "8月", "完成3类设备校验"], ["伙伴试用", "9月", "误报率≤3%"], ["正式发布", "10月15日", "支持5类设备"]]),
            _slide("closing", "2026年10月15日，与首批伙伴一起开启零接触巡检", bullets=["报名参与9月伙伴试用", "预约工厂网络与设备兼容性评估", "获取上线清单与安全边界说明"], callout="星澜 S1 · 现场可验证"),
        ],
        "content_tokens": ["零接触巡检", "工业设备智能巡检终端", "三阶段发布", "伙伴试用", "现场可验证"],
        "numeric_tokens": ["45 分钟", "12 分钟", "99.95%", "8.2%", "2.6%", "2026年10月15日"],
        "preserve_tokens": ["星澜 S1", "2026.10.15"],
        "forbidden_tokens": ["永不宵机", "100%准确", "全行业已验证"],
        "object_mins": {"charts": 1, "tables": 1, "pictures": 1, "shapes": 50},
        "references": ["PPTC/README.md", "PPTC/PPT_test_input/Create_new_slides/session_0.json"],
        "requirements": ["把大纲变成7页可直接用于发布会的故事线", "保留所有给定的时间、可用性与试点数据", "使用品牌配色，但不做无依据的市场承诺"],
    },
    {
        "id": "ppt-003-east-china-retail-weekly",
        "family": "data_to_deck",
        "title": "华东零售周报数据汇报",
        "sector": "零售",
        "occupation": "商业分析师",
        "difficulty": "advanced",
        "estimated": 90,
        "capabilities": ["K2", "K4", "K5", "K7", "K9"],
        "accent": "F39C12",
        "author": "澄日零售（合成）",
        "inputs": [
            {
                "name": "华东零售第32周数据.xlsx",
                "kind": "xlsx",
                "sheets": [
                    {"name": "城市", "rows": [["城市", "销售额(万元)", "环比", "退货率"], ["上海", 186, 0.148, 0.019], ["杭州", 121, 0.117, 0.022], ["南京", 96, 0.082, 0.025], ["苏州", 79, 0.155, 0.018], ["合计", 482, 0.124, 0.021]]},
                    {"name": "渠道", "rows": [["渠道", "销售额(万元)", "毛利率"], ["门店", 278, 0.346], ["小程序", 137, 0.382], ["平台", 67, 0.291]]},
                    {"name": "品类", "rows": [["品类", "销售额(万元)", "贡献增量(万元)"], ["轻食", 148, 19], ["咖啡", 132, 17], ["烘焙", 116, 9], ["周边", 86, 8]]},
                ],
            },
            {"name": "周报口径说明.md", "kind": "text", "content": "# 第32周口径\n\n销售额为含税实收，环比对比第31周。全区销售额482万元，环比+12.4%；上海186万元。全区退货率2.1%。下周资源优先投入苏州新店与上海轻食复购活动。"},
        ],
        "source_markers": ["482", "12.4%", "186", "2.1%", "第32周"],
        "deliverables": [{"name": "华东零售第32周经营简报.pptx", "type": "pptx"}],
        "slides": [
            _slide("cover", "华东零售第32周经营简报", subtitle="含税实收口径 · 对比第31周", callout="W32"),
            _slide("kpi", "482 万元销售额，环比 +12.4%", kpis=[["482万元", "销售额"], ["+12.4%", "周环比"], ["2.1%", "退货率"]], bullets=["四城市全部正增长", "退货率仍在3%内"]),
            _slide("chart", "上海贡献 186 万元，苏州环比最快", chart={"categories": ["上海", "杭州", "南京", "苏州"], "series": [["销售额(万元)", [186, 121, 96, 79]]]}, bullets=["上海占全区38.6%", "苏州环比+15.5%，下周需配合新店补货"]),
            _slide("table", "小程序毛利率 38.2%，但规模仅占 28%", table=[["渠道", "销售额", "毛利率"], ["门店", "278万元", "34.6%"], ["小程序", "137万元", "38.2%"], ["平台", "67万元", "29.1%"]], callout="增量优先放到可控的私域渠道"),
            _slide("chart", "轻食与咖啡贡献 36 万元增量", chart={"categories": ["轻食", "咖啡", "烘焙", "周边"], "series": [["贡献增量(万元)", [19, 17, 9, 8]]]}, bullets=["轻食复购是上海的首要抓手", "周边产品不扩SKU，先清理低周转库存"]),
            _slide("closing", "下周聚焦两件事：苏州新店供给 + 上海轻食复购", bullets=["苏州：将TOP20 SKU安全库存提至1.4倍", "上海：针对30天内轻食购买用户做二次触达", "次周复盘：销售额、毛利率、退货率三指标配对观察"], callout="数据不只是回顾，要对应下一步资源"),
        ],
        "content_tokens": ["含税实收", "苏州新店", "上海轻食复购", "私域渠道", "TOP20 SKU"],
        "numeric_tokens": ["482万元", "+12.4%", "2.1%", "186 万元", "38.2%", "36 万元"],
        "preserve_tokens": ["第32周", "对比第31周"],
        "forbidden_tokens": ["不含税", "同比+12.4%", "退货率1.2%"],
        "object_mins": {"charts": 2, "tables": 1, "pictures": 0, "shapes": 45},
        "references": ["PPT-Eval/task_registry/tasks.json", "Workspace-Bench/data/lite_task_rubrics_pass_rate.csv"],
        "requirements": ["将工作簿数据整理成6页经营简报", "区分销售额、环比、退货率与毛利率口径", "每个分析结论都要对应一个下周动作"],
    },
    {
        "id": "ppt-004-customer-success-template",
        "family": "template_based_generation",
        "title": "客户成功月报套模板生成",
        "sector": "企业软件",
        "occupation": "客户成功经理",
        "difficulty": "advanced",
        "estimated": 90,
        "capabilities": ["K2", "K4", "K6", "K7", "K10"],
        "accent": "27AE60",
        "author": "青岚科技（合成）",
        "inputs": [
            {
                "name": "青岚科技演示模板.pptx",
                "kind": "pptx",
                "accent": "27AE60",
                "footer": "青岚科技｜客户成功部",
                "slides": [
                    _slide("cover", "青岚科技标准封面", subtitle="模板示例：深蓝底 + 绿色强调", callout="Template 2026"),
                    _slide("kpi", "指标页样式", kpis=[["98%", "指标A"], ["8h", "指标B"], ["6/8", "指标C"]]),
                    _slide("closing", "结论页样式", bullets=["一句结论", "一个行动", "一个责任人"], callout="青岚科技｜客户成功部"),
                ],
            },
            {
                "name": "客户成功七月数据.xlsx",
                "kind": "xlsx",
                "sheets": [{"name": "指标", "rows": [["指标", "7月值", "目标", "备注"], ["平台可用性", "98.6%", "99.0%", "华南区域一次故障"], ["平均首响", "7.4h", "8h", "达标"], ["成功计划完成", "6/8", "8/8", "两项延后"], ["高风险客户", 2, 0, "海星物流、东湾零售"]]}],
            },
            {"name": "月报口径.md", "kind": "text", "content": "# 7月客户成功月报\n\n可用性以客户侧实际请求计算；首响为工单创建到首次人工回复。延后的两项成功计划均需在8月20日前关闭。"},
        ],
        "source_markers": ["青岚科技｜客户成功部", "98.6%", "7.4h", "6/8", "8月20日"],
        "deliverables": [{"name": "青岚科技_7月客户成功月报.pptx", "type": "pptx"}],
        "slides": [
            _slide("cover", "7月客户成功月报", subtitle="服务健康度、成功计划与风险客户", callout="青岚科技｜客户成功部"),
            _slide("kpi", "首响达标，可用性与成功计划仍有缺口", kpis=[["98.6%", "平台可用性"], ["7.4h", "平均首响"], ["6/8", "成功计划完成"]]),
            _slide("chart", "可用性低于 99.0% 目标 0.4 个百分点", chart={"categories": ["实际", "目标"], "series": [["可用性", [98.6, 99.0]]]}, bullets=["缺口来自华南区域一次故障", "改进项需与8月稳定性发布配对"]),
            _slide("table", "两个高风险客户需单独建立恢复计划", table=[["客户", "风险信号", "本月动作"], ["海星物流", "接口超时3次", "完成容量评估"], ["东湾零售", "关键用户活跃下降", "开展管理员培训"]]),
            _slide("timeline", "两项延后计划必须在 8月20日前关闭", timeline=[["8月5日", "确认责任人与交付物"], ["8月12日", "客户联合验收"], ["8月20日", "关闭计划与归档证据"]]),
            _slide("closing", "8月管理重点：稳定性、两项计划、两个风险客户", bullets=["可用性回到99.0%以上", "延后计划在8月20日前全部关闭", "高风险客户每周更新恢复信号"], callout="青岚科技｜客户成功部"),
        ],
        "content_tokens": ["客户成功月报", "海星物流", "东湾零售", "华南区域", "恢复计划"],
        "numeric_tokens": ["98.6%", "99.0%", "7.4h", "6/8", "8月20日", "0.4 个百分点"],
        "preserve_tokens": ["青岚科技｜客户成功部", "7月"],
        "forbidden_tokens": ["可用性已达标", "8/8", "无高风险客户"],
        "object_mins": {"charts": 1, "tables": 1, "pictures": 0, "shapes": 45},
        "references": ["PPTC/README.md", "PPTArena/src/evaluation_pairs_refined.json"],
        "requirements": ["使用提供的PPTX作为视觉模板，生成6页月报", "保留深蓝底、绿色强调和部门页脚", "不得把未达标数据包装为已达标"],
    },
    {
        "id": "ppt-005-renewal-review-revision",
        "family": "content_editing",
        "title": "客户续约评审内容修订",
        "sector": "企业服务",
        "occupation": "商务运营经理",
        "difficulty": "advanced",
        "estimated": 75,
        "capabilities": ["K2", "K6", "K8", "K10"],
        "accent": "2F80ED",
        "author": "柏舟云服务（合成）",
        "inputs": [
            {
                "name": "客户续约评审_待修订.pptx",
                "kind": "pptx",
                "accent": "7F8C8D",
                "footer": "柏舟云服务｜内部资料",
                "slides": [
                    _slide("cover", "Q2客户续约评审", subtitle="商务运营会议", callout="待修订"),
                    _slide("kpi", "Q2续约盘点", kpis=[["78%", "续约率"], ["1,260万元", "ARR"], ["5", "高风险客户"]]),
                    _slide("table", "高风险客户", table=[["客户", "责任人", "动作"], ["明辉制造", "陈熙", "待确认"], ["山海药业", "李博", "待确认"]]),
                    _slide("bullets", "续约策略", bullets=["所有客户统一折扣", "风险客户由销售自行跟进"]),
                    _slide("closing", "下一步", bullets=["确认报价", "跟进客户"], callout="柏舟云服务｜内部资料"),
                ],
            },
            {"name": "续约评审修订清单.md", "kind": "text", "content": "# 修订清单\n\n1. 报告期从Q2改为Q3。\n2. 续约率从78%改为84%。\n3. ARR从1,260万元改为1,380万元。\n4. 高风险客户从5家改为3家。\n5. 明辉制造责任人改为王芳，截止日9月12日；山海药业责任人改为赵晨，截止日9月15日。\n6. 保留“柏舟云服务｜内部资料”页脚和5页结构。"},
        ],
        "source_markers": ["Q2", "84%", "1,380万元", "王芳", "柏舟云服务｜内部资料"],
        "deliverables": [{"name": "Q3客户续约评审_已修订.pptx", "type": "pptx"}],
        "slides": [
            _slide("cover", "Q3客户续约评审", subtitle="商务运营会议", callout="已修订"),
            _slide("kpi", "Q3续约盘点：续约率提升6个百分点", kpis=[["84%", "续约率"], ["1,380万元", "ARR"], ["3家", "高风险客户"]]),
            _slide("table", "两家重点客户已明确责任人和截止日", table=[["客户", "责任人", "截止日"], ["明辉制造", "王芳", "9月12日"], ["山海药业", "赵晨", "9月15日"]]),
            _slide("bullets", "续约策略从统一折扣改为分层应对", bullets=["绿色：按标准节奏推进续约", "黄色：提前30天复盘价值证据", "红色：责任人与管理层联合恢复"]),
            _slide("closing", "下一步：用截止日管理两个重点续约", bullets=["王芳：9月12日前完成明辉制造价值复盘", "赵晨：9月15日前完成山海药业方案确认", "周会按续约概率与ARR风险双维跟踪"], callout="柏舟云服务｜内部资料"),
        ],
        "content_tokens": ["Q3客户续约评审", "明辉制造", "山海药业", "王芳", "赵晨", "分层应对"],
        "numeric_tokens": ["84%", "1,380万元", "3家", "9月12日", "9月15日", "6个百分点"],
        "preserve_tokens": ["柏舟云服务｜内部资料", "商务运营会议"],
        "forbidden_tokens": ["Q2客户续约评审", "78%", "1,260万元", "陈熙", "李博"],
        "object_mins": {"charts": 0, "tables": 1, "pictures": 0, "shapes": 38},
        "references": ["PPTArena/src/evaluation_pairs_refined.json", "PPT-Eval/task_registry/rubrics/3-002.json"],
        "requirements": ["严格按修订清单更新原稿，保留5页与页脚", "修正期间、续约率、ARR、风险客户数和责任人", "除明确修订项外不得带入旧值或旧人名"],
    },
    {
        "id": "ppt-006-warehouse-safety-layout",
        "family": "visual_editing",
        "title": "仓储安全培训视觉修订",
        "sector": "物流",
        "occupation": "安全培训专员",
        "difficulty": "intermediate",
        "estimated": 75,
        "capabilities": ["K6", "K7", "K8", "K10"],
        "accent": "E15B2A",
        "author": "速达仓配（合成）",
        "inputs": [
            {
                "name": "仓储安全培训_待排版.pptx",
                "kind": "pptx",
                "accent": "D35400",
                "footer": "速达仓配｜安全培训",
                "slides": [
                    _slide("cover", "仓储安全必修课", subtitle="新员工入场培训", callout="原稿"),
                    _slide("bullets", "三条红线", bullets=["人车分流", "严禁跨越输送线", "货架变形立即上报"]),
                    _slide("timeline", "叉车作业前检查", timeline=[["1", "轮胎与货叉"], ["2", "制动与转向"], ["3", "灯光与蜂鸣器"]]),
                    _slide("table", "应急联系", table=[["场景", "动作"], ["火警", "报警并沿绿色通道撤离"], ["人员受伤", "保护现场并联系值班长"]]),
                    _slide("closing", "确认与测验", bullets=["完成现场签到", "测试达80分后上岗"], callout="速达仓配｜安全培训"),
                ],
            },
            {"name": "仓储培训视觉规范.md", "kind": "text", "content": "# 视觉规范\n\n- 比例16:9，标题不小于28pt，正文不小于18pt。\n- 主色#1F2937，警示色#E15B2A，空白底色#F7F8FA。\n- 每页仅保留一个主结论，三条红线必须可在3秒内读完。\n- 保留“速达仓配｜安全培训”页脚。\n- 不得改变原稿的安全规则与80分上岗门槛。"},
        ],
        "source_markers": ["人车分流", "严禁跨越输送线", "80分", "#E15B2A", "速达仓配｜安全培训"],
        "deliverables": [{"name": "仓储安全培训_已优化.pptx", "type": "pptx"}],
        "slides": [
            _slide("cover", "仓储安全必修课", subtitle="新员工入场培训｜先识别风险，再开始作业", callout="安全第一"),
            _slide("kpi", "三条红线：任何一条都不能侵犯", kpis=[["01", "人车分流"], ["02", "严禁跨越输送线"], ["03", "货架变形立即上报"]]),
            _slide("timeline", "叉车作业前：按同一方向完成三步检查", timeline=[["1 车前", "轮胎与货叉"], ["2 驾驶位", "制动与转向"], ["3 车后", "灯光与蜂鸣器"]], callout="检查异常 -> 停止作业 -> 上报值班长"),
            _slide("table", "紧急情况：先保护人，再保护现场", table=[["场景", "第一动作", "后续"], ["火警", "报警", "沿绿色通道撤离"], ["人员受伤", "停止作业", "保护现场并联系值班长"]]),
            _slide("closing", "测试达 80 分后上岗", bullets=["在场完成签到与设备点检演示", "不足80分者当日补训并重测", "安全疑问先问值班长，不凭经验冒险作业"], callout="速达仓配｜安全培训"),
        ],
        "content_tokens": ["人车分流", "严禁跨越输送线", "货架变形立即上报", "绿色通道", "值班长"],
        "numeric_tokens": ["80 分", "01", "02", "03"],
        "preserve_tokens": ["速达仓配｜安全培训", "新员工入场培训"],
        "forbidden_tokens": ["跨越输送线是允许的", "60分后上岗", "异常时继续作业"],
        "object_mins": {"charts": 0, "tables": 1, "pictures": 0, "shapes": 42},
        "references": ["PPTArena/src/evaluation_pairs_refined.json", "PPTArena/GroundTruth/FixImage&TextOverFlow_GroundTruthB.pptx"],
        "requirements": ["只优化视觉层级、对齐、字号和配色，不改变安全语义", "保留5页结构、三条红线和80分上岗门槛", "页面不得有遮挡、溢出或低对比度正文"],
    },
    {
        "id": "ppt-007-quarterly-ops-polish",
        "family": "deck_polishing",
        "title": "季度经营会演示稿润色",
        "sector": "生活服务",
        "occupation": "经营分析经理",
        "difficulty": "advanced",
        "estimated": 90,
        "capabilities": ["K4", "K6", "K7", "K8", "K10"],
        "accent": "0C7C86",
        "author": "柚子生活（合成）",
        "inputs": [
            {
                "name": "二季度经营会_初稿.pptx",
                "kind": "pptx",
                "accent": "95A5A6",
                "footer": "柚子生活｜内部经营会",
                "slides": [
                    _slide("cover", "2026年二季度经营会", subtitle="业务回顾和第三季度计划"),
                    _slide("bullets", "总体结果", bullets=["收入2,480万元", "贡献毛利率41.2%", "付费会员15.6万"]),
                    _slide("chart", "月度收入", chart={"categories": ["4月", "5月", "6月"], "series": [["收入(万元)", [760, 810, 910]]]}),
                    _slide("table", "门店表现", table=[["城市", "收入", "利润率"], ["深圳", "860万元", "13.8%"], ["广州", "740万元", "12.6%"], ["厦门", "420万元", "11.9%"]]),
                    _slide("bullets", "主要问题", bullets=["拉新成本上升", "厦门门店人效偏低", "会员复购率下降1.8个百分点"]),
                    _slide("closing", "Q3行动计划", bullets=["优化渠道", "提升人效", "恢复复购"]),
                ],
            },
            {"name": "经营会润色要求.md", "kind": "text", "content": "# 润色要求\n\n保留6页和所有业务数据。把页标题改成“结论句”；全稿使用#17324D主色、#0C7C86强调色；将Q3计划写成有责任人和量化目标的动作。保留“柚子生活｜内部经营会”页脚。"},
        ],
        "source_markers": ["2,480万元", "41.2%", "15.6万", "1.8个百分点", "柚子生活｜内部经营会"],
        "deliverables": [{"name": "2026Q2经营会_润色版.pptx", "type": "pptx"}],
        "slides": [
            _slide("cover", "2026 Q2：收入增长，但拉新与复购同时承压", subtitle="柚子生活季度经营会", callout="Q2 REVIEW"),
            _slide("kpi", "2,480 万元收入达成，毛利率保持 41.2%", kpis=[["2,480万元", "收入"], ["41.2%", "贡献毛利率"], ["15.6万", "付费会员"]]),
            _slide("chart", "6月收入 910 万元，连续三月增长", chart={"categories": ["4月", "5月", "6月"], "series": [["收入(万元)", [760, 810, 910]]]}, bullets=["增长主要由6月会员日与门店联动带动", "下季度不将单次促销外推为自然增长"]),
            _slide("table", "深圳贡献最高，厦门利润率仅 11.9%", table=[["城市", "收入", "利润率", "判断"], ["深圳", "860万元", "13.8%", "保持"], ["广州", "740万元", "12.6%", "提效"], ["厦门", "420万元", "11.9%", "重点改善"]]),
            _slide("bullets", "三个问题都已找到可量化的管理抓手", bullets=["拉新：将付费渠道CAC上限设为68元", "人效：厦门排班按小时客流重算", "复购：针对60天未访会员做分层召回"], callout="复购率较上季度下降1.8个百分点"),
            _slide("closing", "Q3：责任到人，指标到周", bullets=["市场负责人林晴：将付费渠道CAC降至68元以内", "区域负责人周扬：将厦门利润率提至12.5%", "会员负责人许文：在9月末前恢复1.0个百分点复购率"], callout="柚子生活｜内部经营会"),
        ],
        "content_tokens": ["柚子生活季度经营会", "厦门", "分层召回", "责任到人", "指标到周"],
        "numeric_tokens": ["2,480万元", "41.2%", "15.6万", "910 万元", "11.9%", "1.8个百分点", "68元", "12.5%"],
        "preserve_tokens": ["柚子生活｜内部经营会", "Q2 REVIEW"],
        "forbidden_tokens": ["数据待定", "责任人待定", "保证恢复"],
        "object_mins": {"charts": 1, "tables": 1, "pictures": 0, "shapes": 46},
        "references": ["PPTArena/GroundTruth/AestheticMakeover_GroundTruthA.pptx", "PPT-Eval/proposed_tasks/presentation-tips.proposed_tasks.json"],
        "requirements": ["保留6页和原始业务数据，将标题改为可单独理解的结论句", "统一配色、栅格、字号、图表与表格样式", "Q3行动必须具备责任人和可验证目标"],
    },
    {
        "id": "ppt-008-product-lines-consolidation",
        "family": "deck_consolidation",
        "title": "产品线经营稿合并",
        "sector": "智能硬件",
        "occupation": "产品运营总监",
        "difficulty": "expert",
        "estimated": 120,
        "capabilities": ["K2", "K4", "K6", "K8", "K10", "K13"],
        "accent": "9B51E0",
        "author": "云纳智造（合成）",
        "inputs": [
            {
                "name": "产品线A月报.pptx",
                "kind": "pptx",
                "accent": "2D9CDB",
                "footer": "云纳智造｜产品线A",
                "slides": [
                    _slide("cover", "产品线A｜7月经营月报", subtitle="工业相机"),
                    _slide("kpi", "A线结果", kpis=[["860万元", "收入"], ["38%", "毛利率"], ["93%", "交付及时率"]]),
                    _slide("bullets", "A线机会", bullets=["新能源客户追加订单", "A3型号9月进入量产"]),
                    _slide("closing", "A线请求", bullets=["追加采购额80万元"]),
                ],
            },
            {
                "name": "产品线B月报.pptx",
                "kind": "pptx",
                "accent": "EB5757",
                "footer": "云纳智造｜产品线B",
                "slides": [
                    _slide("cover", "产品线B｜7月经营月报", subtitle="边缘网关"),
                    _slide("kpi", "B线结果", kpis=[["640万元", "收入"], ["31%", "毛利率"], ["87%", "交付及时率"]]),
                    _slide("bullets", "B线风险", bullets=["B2芯片到货延迟", "4个客户需重排交付"]),
                    _slide("closing", "B线请求", bullets=["启用备选芯片方案"]),
                ],
            },
            {"name": "产品线合并口径.xlsx", "kind": "xlsx", "sheets": [{"name": "口径", "rows": [["产品线", "收入(万元)", "毛利率", "交付及时率", "下月目标"], ["A", 860, 0.38, 0.93, 920], ["B", 640, 0.31, 0.87, 720], ["合计", 1500, 0.35, 0.905, 1640]]}]},
        ],
        "source_markers": ["860万元", "640万元", "1500", "0.905", "B2芯片"],
        "deliverables": [{"name": "云纳智造_7月产品线联合经营报告.pptx", "type": "pptx"}],
        "slides": [
            _slide("cover", "7月产品线联合经营报告", subtitle="工业相机A线 + 边缘网关B线", callout="联合决策版"),
            _slide("kpi", "两条产品线合计收入 1,500 万元", kpis=[["1,500万元", "合计收入"], ["35.0%", "综合毛利率"], ["90.5%", "交付及时率"]]),
            _slide("chart", "A线贡献 57% 收入，B线是交付短板", chart={"categories": ["A线", "B线"], "series": [["收入(万元)", [860, 640]]]}, bullets=["A线收入860万元、毛利率38%", "B线收入640万元、交付及时率87%"]),
            _slide("table", "A线抓增量，B线先恢复交付", table=[["产品线", "优势/问题", "八月动作"], ["A线", "新能源追加订单", "A3量产前置备料"], ["B线", "B2芯片到货延迟", "验证备选芯片并重排4客户"]]),
            _slide("timeline", "共用供应链节点，但不混合产品责任", timeline=[["8月5日", "B2备选芯片验证"], ["8月12日", "A3量产备料锁定"], ["8月20日", "四个客户交付日确认"]]),
            _slide("chart", "8月目标 1,640 万元，增量主要来自B线恢复", chart={"categories": ["A线", "B线"], "series": [["7月", [860, 640]], ["8月目标", [920, 720]]]}, bullets=["目标不包含未签约的潜在项目", "B线目标以备选芯片验证通过为前提"]),
            _slide("table", "两项资源请求需分别决策", table=[["请求", "额度/边界", "决策条件"], ["A线追加采购", "80万元", "与A3已确认订单绑定"], ["B线备选芯片", "小批量试采", "通过稳定性与兼容性验证"]]),
            _slide("closing", "建议：批准A线有条件备料，B线先验证再放量", bullets=["A线：采购额与已确认订单配对", "B线：不在验证前扩大备选芯片使用", "联合周会同时跟踪收入、毛利率和交付及时率"], callout="云纳智造｜内部决策资料"),
        ],
        "content_tokens": ["工业相机A线", "边缘网关B线", "B2芯片", "A3量产", "有条件备料", "先验证再放量"],
        "numeric_tokens": ["1,500万元", "35.0%", "90.5%", "860万元", "640万元", "1,640 万元", "80万元"],
        "preserve_tokens": ["联合决策版", "云纳智造｜内部决策资料"],
        "forbidden_tokens": ["A线与B线数据完全可比", "B2已解决", "无条件批准"],
        "object_mins": {"charts": 2, "tables": 2, "pictures": 0, "shapes": 60},
        "references": ["PPTArena/GroundTruth/CrossSlideDataConsolidation_GroundTruthA.pptx", "PPT-Eval/task_registry/tasks.json"],
        "requirements": ["将两份月报合并成8页管理层稿，去掉重复封面和冗余结论", "按统一口径展示合计，同时保留A/B线独立责任", "资源请求必须有额度、条件与前置风险"],
    },
    {
        "id": "ppt-009-architecture-executive-adaptation",
        "family": "audience_adaptation",
        "title": "技术架构稿管理层改写",
        "sector": "互联网平台",
        "occupation": "技术项目经理",
        "difficulty": "expert",
        "estimated": 105,
        "capabilities": ["K1", "K3", "K6", "K8", "K10"],
        "accent": "00A8A8",
        "author": "曙光交易平台（合成）",
        "inputs": [
            {
                "name": "交易平台架构评审_技术版.pptx",
                "kind": "pptx",
                "accent": "34495E",
                "footer": "曙光交易平台｜架构评审",
                "slides": [
                    _slide("cover", "交易平台高可用改造", subtitle="技术架构评审"),
                    _slide("bullets", "现状架构", bullets=["Kafka ISR报警", "Kubernetes HPA扩容滞后", "P99 GC pause 380ms"]),
                    _slide("chart", "故障恢复", chart={"categories": ["改造前", "试点后"], "series": [["MTTR(分钟)", [47, 18]]]}),
                    _slide("table", "改造项", table=[["组件", "方案"], ["消息队列", "ISR分层告警"], ["容器", "HPA预热"], ["JVM", "GC参数调优"]]),
                    _slide("kpi", "容量与成本", kpis=[["99.95%", "SLO"], ["320万元", "年化节省"], ["6周", "上线周期"]]),
                    _slide("bullets", "风险", bullets=["切流期间需安排回退窗口", "试点结果尚未覆盖大促峰值"]),
                    _slide("closing", "实施计划", bullets=["灰度、扩大、全量"]),
                ],
            },
            {"name": "管理层沟通要求.md", "kind": "text", "content": "# 管理层沟通\n\n受众为CEO、CFO和业务副总。保留可用性99.95%、故障恢复时间从47分钟降至18分钟、年化节省320万元、6周上线周期与“尚未覆盖大促峰值”风险。把Kafka ISR、Kubernetes HPA、P99 GC pause等实现术语改写为业务影响，不得伪装为零风险。"},
        ],
        "source_markers": ["99.95%", "47分钟", "18分钟", "320万元", "Kafka ISR"],
        "deliverables": [{"name": "交易平台高可用改造_管理层版.pptx", "type": "pptx"}],
        "slides": [
            _slide("cover", "交易平台高可用改造｜管理层决策稿", subtitle="用更快恢复与更低运行成本支撑业务增长", callout="投资评审"),
            _slide("kpi", "改造后，故障恢复时间从 47 降至 18 分钟", kpis=[["99.95%", "目标可用性"], ["18分钟", "试点恢复时间"], ["320万元", "年化节省"]]),
            _slide("chart", "恢复时间缩短 62%，业务暴露窗口明显下降", chart={"categories": ["改造前", "试点后"], "series": [["故障恢复时间(分钟)", [47, 18]]]}, bullets=["意味着订单中断时间更短", "试点结果不等于大促峰值保证"]),
            _slide("table", "投资对应三个可验证的业务收益", table=[["投入方向", "业务收益", "验证指标"], ["更早识别风险", "减少突发中断", "事前告警覆盖率"], ["需求到来前扩容", "降低高峰排队", "订单延迟率"], ["缩短恢复路径", "缩短业务中断", "恢复时间"]]),
            _slide("timeline", "6 周分三阶段上线，每阶段都保留回退能力", timeline=[["1–2周", "5%流量验证"], ["3–4周", "30%流量 + 峰值演练"], ["5–6周", "分时段全量切换"]], callout="任一门槛未达成 -> 暂停扩大"),
            _slide("closing", "建议批准：以 99.95% 和大促演练作为上线门槛", bullets=["批准6周实施窗口与配套资源", "要求第4周前完成大促峰值演练", "财务收益按实际节省复核，不将试点估算当作已实现收益"], callout="曙光交易平台｜管理层决策"),
        ],
        "content_tokens": ["管理层决策稿", "业务暴露窗口", "大促峰值演练", "回退能力", "实际节省复核"],
        "numeric_tokens": ["99.95%", "47 ", "18 分钟", "62%", "320万元", "6 周", "5%", "30%"],
        "preserve_tokens": ["曙光交易平台", "投资评审"],
        "forbidden_tokens": ["Kafka ISR", "Kubernetes HPA", "P99 GC pause", "零风险", "已实现320万元"],
        "object_mins": {"charts": 1, "tables": 1, "pictures": 0, "shapes": 45},
        "references": ["PPT-Eval/proposed_tasks/Qatargas GS new Customer Service.proposed_tasks.json", "PPTArena/GroundTruth/ShortenTextOnSlides_GroundTruthA.pptx"],
        "requirements": ["把7页技术评审改写为6页CEO/CFO可决策的稿件", "将实现术语翻译为业务影响，保留所有边界与风险", "明确投资收益、六周节奏、门槛和回退机制"],
    },
    {
        "id": "ppt-010-apac-launch-localization",
        "family": "localization",
        "title": "APAC 发布计划中文本地化",
        "sector": "云计算",
        "occupation": "区域市场经理",
        "difficulty": "advanced",
        "estimated": 90,
        "capabilities": ["K3", "K6", "K8", "K10"],
        "accent": "D94F70",
        "author": "Aster Cloud（合成）",
        "inputs": [
            {
                "name": "APAC Launch Plan_EN.pptx",
                "kind": "pptx",
                "accent": "D94F70",
                "footer": "Aster Cloud | Confidential",
                "slides": [
                    _slide("cover", "APAC Launch Plan", subtitle="Aster Cloud Edge Analytics", callout="Go-live: 15 October 2026"),
                    _slide("kpi", "Launch Overview", kpis=[["3", "Launch markets"], ["12", "Channel partners"], ["US$1.8M", "Pipeline target"]]),
                    _slide("table", "Market Readiness", table=[["Market", "Readiness", "Owner"], ["Mainland China", "Ready", "Regional Marketing"], ["Singapore", "Ready", "Partner Team"], ["Japan", "At risk", "Product Marketing"]]),
                    _slide("timeline", "Go-to-market Timeline", timeline=[["August", "Partner enablement"], ["September", "Customer preview"], ["15 October", "General availability"]]),
                    _slide("bullets", "Risks and Mitigations", bullets=["Japan documentation review may delay launch", "Chinese terminology must match the approved glossary", "No customer data may leave the deployment region"]),
                    _slide("closing", "Decision Required", bullets=["Approve US$240K launch budget", "Confirm 15 October 2026 go-live"], callout="Aster Cloud | Confidential"),
                ],
            },
            {"name": "APAC本地化术语表.xlsx", "kind": "xlsx", "sheets": [{"name": "术语", "rows": [["English", "简体中文", "备注"], ["Edge Analytics", "边缘分析", "产品名保留Aster Cloud"], ["Go-live", "正式上线", "日期写为2026年10月15日"], ["General availability", "全面可用", ""], ["Pipeline", "销售线索金额", "US$换算为美元，不进行汇率换算"], ["Confidential", "机密", "页脚用语"]]}]},
        ],
        "source_markers": ["APAC Launch Plan", "15 October 2026", "US$1.8M", "US$240K", "边缘分析"],
        "deliverables": [{"name": "AsterCloud_APAC发布计划_简体中文.pptx", "type": "pptx"}],
        "slides": [
            _slide("cover", "APAC 发布计划", subtitle="Aster Cloud 边缘分析", callout="正式上线：2026年10月15日"),
            _slide("kpi", "发布概览：3个市场、12家渠道伙伴", kpis=[["3个", "首发市场"], ["12家", "渠道伙伴"], ["180万美元", "销售线索金额目标"]]),
            _slide("table", "中国大陆与新加坡就绪，日本文档审核有风险", table=[["市场", "准备状态", "责任团队"], ["中国大陆", "已就绪", "区域市场"], ["新加坡", "已就绪", "伙伴团队"], ["日本", "有风险", "产品市场"]]),
            _slide("timeline", "从伙伴赋能到全面可用，三个节点不变", timeline=[["8月", "伙伴赋能"], ["9月", "客户预览"], ["10月15日", "全面可用"]]),
            _slide("bullets", "三项本地化风险都已配对应对动作", bullets=["日本文档审核：设置9月25日冻结点", "中文术语：只使用已批准术语表", "数据合规：客户数据不得离开部署区域"]),
            _slide("closing", "待决策：批准 24 万美元预算并确认上线日", bullets=["发布预算：24万美元", "正式上线：2026年10月15日", "日本市场以文档审核通过为发布前置条件"], callout="Aster Cloud｜机密"),
        ],
        "content_tokens": ["Aster Cloud 边缘分析", "中国大陆", "新加坡", "日本文档审核", "全面可用", "机密"],
        "numeric_tokens": ["2026年10月15日", "3个", "12家", "180万美元", "24万美元", "9月25日"],
        "preserve_tokens": ["Aster Cloud", "APAC", "Aster Cloud｜机密"],
        "forbidden_tokens": ["Launch Overview", "Market Readiness", "Go-to-market Timeline", "Decision Required", "Aster Cloud | Confidential"],
        "object_mins": {"charts": 0, "tables": 1, "pictures": 0, "shapes": 45},
        "references": ["PPTArena/src/evaluation_pairs_refined.json", "PPTArena/GroundTruth/TranslateArabicLTR_GroundTruthB.pptx"],
        "requirements": ["将六页英文稿本地化为简体中文，保持页数、数字、市场顺序与版式", "严格使用术语表，Aster Cloud与APAC保持原文", "将日期与美元金额按中文习惯表达，不进行汇率换算"],
    },
    {
        "id": "ppt-011-volunteer-conference-package",
        "family": "presentation_package",
        "title": "志愿者大会演示交付包",
        "sector": "公益",
        "occupation": "会务项目经理",
        "difficulty": "expert",
        "estimated": 120,
        "capabilities": ["K1", "K2", "K4", "K7", "K9", "K12"],
        "accent": "E67E22",
        "author": "微光公益联盟（合成）",
        "inputs": [
            {"name": "年度志愿者大会素材.docx", "kind": "docx", "title": "2026年微光志愿者大会素材", "paragraphs": ["主题：让每一次善意都有回声。", "年度成果：1,280名活跃志愿者，累计服务32,640小时，覆盖21个社区，受益8,900人次。", "年度故事：夜间陪诊计划由80人扩展至230人，但仍需完善排班与志愿者保护。", "明年重点：新增6个社区服务点，将核心志愿者留存率由72%提升至80%。"], "table": [["环节", "时间", "主持/负责"], ["年度回顾", "09:30", "项目组"], ["志愿者故事", "10:10", "社区伙伴"], ["明年计划", "11:00", "秘书处"]]},
            {"name": "参会数据.xlsx", "kind": "xlsx", "sheets": [{"name": "参会", "rows": [["类别", "人数", "占比"], ["志愿者", 286, 0.715], ["社区伙伴", 72, 0.18], ["捐赠方", 24, 0.06], ["工作人员", 18, 0.045], ["合计", 400, 1.0]]}]},
            {"name": "会务信息.md", "kind": "text", "content": "# 会务\n\n日期：2026年12月5日\n地点：湖滨市民中心二层大礼堂\n入场：09:00\n正式开始：09:30\n无障碍入口：东门\n紧急联系：会务台（不公开个人手机号）"},
            {"name": "大会主视觉.png", "kind": "png", "palette": ["FFF2D8", "E67E22", "2E4057"]},
        ],
        "source_markers": ["1,280", "32,640", "8,900", "2026年12月5日", "400"],
        "deliverables": [
            {"name": "2026微光志愿者大会.pptx", "type": "pptx"},
            {"name": "2026微光志愿者大会_参会手册.pdf", "type": "pdf"},
            {"name": "2026微光志愿者大会_讲者手册.docx", "type": "docx"},
        ],
        "package_tokens": ["让每一次善意都有回声", "2026年12月5日", "湖滨市民中心", "09:30"],
        "slides": [
            _slide("cover", "让每一次善意都有回声", subtitle="2026微光志愿者大会", callout="2026年12月5日"),
            _slide("kpi", "1,280 名志愿者，用 32,640 小时服务 21 个社区", kpis=[["1,280名", "活跃志愿者"], ["32,640小时", "志愿服务"], ["8,900人次", "受益者"]]),
            _slide("chart", "400 位参会者中，志愿者占 71.5%", chart={"categories": ["志愿者", "社区伙伴", "捐赠方", "工作人员"], "series": [["人数", [286, 72, 24, 18]]]}, bullets=["设置志愿者快速签到通道", "社区伙伴与捐赠方席位相邻，便于交流"]),
            _slide("table", "三个核心环节在 11:30 前完成", table=[["时间", "环节", "负责"], ["09:30", "年度回顾", "项目组"], ["10:10", "志愿者故事", "社区伙伴"], ["11:00", "明年计划", "秘书处"]]),
            _slide("bullets", "夜间陪诊从 80 人扩展至 230 人", bullets=["增长来自社区转介与志愿者口碑", "扩展后的首要问题是排班可持续性", "明年将志愿者保护与心理支持纳入服务标准"]),
            _slide("timeline", "2027：新增 6 个服务点，将核心留存率提升至 80%", timeline=[["Q1", "完成需求地图与合作社区签约"], ["Q2", "启动3个服务点"], ["Q3–Q4", "完成其余3个点并复盘留存"]]),
            _slide("closing", "感谢每一位让善意产生回声的人", bullets=["地点：湖滨市民中心二层大礼堂", "入场：09:00，正式开始：09:30", "无障碍入口：东门；紧急情况联系会务台"], callout="微光公益联盟｜让善意被看见"),
        ],
        "content_tokens": ["让每一次善意都有回声", "夜间陪诊", "志愿者保护", "湖滨市民中心二层大礼堂", "无障碍入口"],
        "numeric_tokens": ["1,280名", "32,640小时", "8,900人次", "400 ", "71.5%", "80 人", "230 人", "80%"],
        "preserve_tokens": ["2026微光志愿者大会", "2026年12月5日"],
        "forbidden_tokens": ["个人手机号", "真实住址", "身份证号"],
        "object_mins": {"charts": 1, "tables": 1, "pictures": 1, "shapes": 52},
        "references": ["OmegaUse-OfficeVal/data/rubrics-en/officeval_004.json", "Workspace-Bench/README.md", "PPT-Eval/README.md"],
        "requirements": ["交付PPTX主稿、PDF参会手册和DOCX讲者手册三个文件", "三个文件的主题、日期、地点、时间与核心数据必须一致", "不得在任何交付物中暴露志愿者个人手机或地址"],
    },
    {
        "id": "ppt-012-board-deck-audit",
        "family": "deck_audit",
        "title": "董事会经营简报审计与更正",
        "sector": "新能源",
        "occupation": "董事会办公室分析师",
        "difficulty": "expert",
        "estimated": 120,
        "capabilities": ["K2", "K4", "K6", "K8", "K11", "K13"],
        "accent": "C7A34B",
        "author": "屿光能源（合成）",
        "inputs": [
            {
                "name": "董事会经营简报_待审计.pptx",
                "kind": "pptx",
                "accent": "C7A34B",
                "footer": "屿光能源｜董事会资料｜未审计",
                "slides": [
                    _slide("cover", "2026上半年经营简报", subtitle="董事会资料", callout="未审计"),
                    _slide("kpi", "核心业绩", kpis=[["12.48亿元", "营收"], ["19.8%", "同比"], ["34.6%", "毛利率"]]),
                    _slide("chart", "季度营收", chart={"categories": ["Q1", "Q2"], "series": [["营收(亿元)", [5.72, 7.12]]]}),
                    _slide("table", "业务板块", table=[["板块", "营收", "同比"], ["储能系统", "7.34亿元", "+24.1%"], ["运维服务", "3.02亿元", "+11.5%"], ["其他", "2.48亿元", "+8.0%"]]),
                    _slide("kpi", "现金流", kpis=[["2.61亿元", "经营现金流"], ["3.4亿元", "期末现金"], ["68天", "应收周转"]]),
                    _slide("closing", "建议", bullets=["继续扩大储能产能"], callout="屿光能源｜董事会资料｜未审计"),
                ],
            },
            {"name": "董事会数据基准.xlsx", "kind": "xlsx", "sheets": [{"name": "经审核口径", "rows": [["指标", "正确值", "单位", "依据"], ["上半年营收", 12.84, "亿元", "Q1 5.72 + Q2 7.12"], ["营收同比", 0.189, "%", "财务复核"], ["综合毛利率", 0.364, "%", "财务复核"], ["经营现金流", 2.16, "亿元", "现金流量表"], ["储能系统营收", 7.34, "亿元", "分部报表"], ["运维服务营收", 3.02, "亿元", "分部报表"], ["其他营收", 2.48, "亿元", "分部报表"]]}]},
            {"name": "董事会品牌与审计规范.md", "kind": "text", "content": "# 规范\n\n主色#132238，强调色#C7A34B。所有金额统一使用亿元，百分比保留一位小数。修正后页脚必须为“屿光能源｜董事会资料｜已核验”。最后一页要记录更正的12.84亿元、18.9%、36.4%和2.16亿元，并说明本文件为经营简报核验，不代替法定审计。"},
        ],
        "source_markers": ["12.84", "18.9%", "36.4%", "2.16", "已核验"],
        "deliverables": [{"name": "2026H1董事会经营简报_已核验.pptx", "type": "pptx"}],
        "slides": [
            _slide("cover", "2026上半年经营简报", subtitle="董事会资料｜核心指标已与数据基准表配对", callout="已核验"),
            _slide("kpi", "营收 12.84 亿元，同比 +18.9%", kpis=[["12.84亿元", "上半年营收"], ["+18.9%", "营收同比"], ["36.4%", "综合毛利率"]]),
            _slide("chart", "Q2 营收 7.12 亿元，高于 Q1 的 5.72 亿元", chart={"categories": ["Q1", "Q2"], "series": [["营收(亿元)", [5.72, 7.12]]]}, bullets=["上半年合计为12.84亿元", "季度数据与年中合计勾稽一致"]),
            _slide("table", "三个板块营收合计 12.84 亿元", table=[["板块", "营收", "同比"], ["储能系统", "7.34亿元", "+24.1%"], ["运维服务", "3.02亿元", "+11.5%"], ["其他", "2.48亿元", "+8.0%"]]),
            _slide("kpi", "经营现金流 2.16 亿元，需与利润增长配对观察", kpis=[["2.16亿元", "经营现金流"], ["3.40亿元", "期末现金"], ["68天", "应收周转"]], bullets=["现金流值来自现金流量表", "后续应按大客户交付节奏管理回款"]),
            _slide("table", "核验结果：四项核心数字已更正", table=[["指标", "待审计稿", "已核验值"], ["营收", "12.48亿元", "12.84亿元"], ["同比", "19.8%", "18.9%"], ["毛利率", "34.6%", "36.4%"], ["经营现金流", "2.61亿元", "2.16亿元"]], callout="本文件为经营简报核验，不代替法定审计"),
        ],
        "content_tokens": ["董事会资料", "核心指标已与数据基准表配对", "储能系统", "运维服务", "不代替法定审计"],
        "numeric_tokens": ["12.84亿元", "+18.9%", "36.4%", "2.16亿元", "7.12 亿元", "5.72 亿元", "3.40亿元", "68天"],
        "preserve_tokens": ["屿光能源｜董事会资料｜已核验", "已核验"],
        "forbidden_tokens": ["核心业绩12.48亿元", "同比19.8%", "毛利率34.6%", "经营现金流2.61亿元", "未审计"],
        "object_mins": {"charts": 1, "tables": 2, "pictures": 0, "shapes": 48},
        "references": ["PPTArena/src/evaluation_pairs_refined.json", "PPT-Eval/task_registry/tasks.json", "OmegaUse-OfficeVal/data/rubrics-en/officeval_003.json"],
        "requirements": ["以数据基准为唯一数值口径，审计并更正六页董事会稿", "保留原稿的业务板块和季度数据，确保勾稽一致", "最后一页列出四项更正且声明不代替法定审计"],
    },
]


def _hex_color(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def _add_text(
    slide: Any,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    color: str = "243447",
    bold: bool = False,
    align: int = PP_ALIGN.LEFT,
    font: str = "Noto Sans CJK SC",
) -> Any:
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.06)
    frame.margin_right = Inches(0.06)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.03)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = PptPt(size)
    run.font.bold = bold
    run.font.color.rgb = _hex_color(color)
    return shape


def _add_rect(
    slide: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: str,
    *,
    radius: bool = True,
    line: str | None = None,
) -> Any:
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _hex_color(fill)
    shape.line.color.rgb = _hex_color(line or fill)
    return shape


def _style_chart(chart: Any, accent: str) -> None:
    chart.has_legend = len(chart.series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.tick_labels.font.size = PptPt(10)
    chart.category_axis.tick_labels.font.size = PptPt(11)
    for series_index, series in enumerate(chart.series):
        colors = [accent, "62B6CB", "F4A261", "7F8C8D"]
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = _hex_color(colors[series_index % len(colors)])


def _render_deck(path: Path, deck: dict[str, Any], input_dir: Path | None = None) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = deck["title"]
    prs.core_properties.subject = deck.get("family", "synthetic-office-task")
    prs.core_properties.author = deck.get("author", "CoworkEvoBench")
    prs.core_properties.keywords = "CoworkEvoBench, synthetic, PowerPoint"
    accent = deck["accent"]
    hero_image = None
    if input_dir:
        for candidate in sorted(input_dir.glob("*.png")):
            hero_image = candidate
            break

    for index, spec in enumerate(deck["slides"]):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        kind = spec["kind"]
        dark = kind in {"cover", "closing"}
        bg = "132238" if dark else ("F3F6F9" if index % 2 else "FFFFFF")
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = _hex_color(bg)

        if kind == "cover":
            _add_rect(slide, 0, 0, 0.22, 7.5, accent, radius=False)
            _add_rect(slide, 0.8, 0.78, 1.25, 0.13, accent, radius=False)
            if hero_image and hero_image.exists():
                slide.shapes.add_picture(str(hero_image), Inches(8.7), Inches(0.75), width=Inches(3.9), height=Inches(4.9))
                _add_rect(slide, 8.7, 5.75, 3.9, 0.12, accent, radius=False)
                title_width = 7.2
            else:
                for n in range(5):
                    _add_rect(slide, 9.1 + n * 0.32, 1.2 + n * 0.55, 2.5 - n * 0.18, 0.18, accent, radius=True)
                title_width = 8.0
            _add_text(slide, spec["title"], 0.82, 1.18, title_width, 2.2, size=30, color="FFFFFF", bold=True)
            if spec.get("subtitle"):
                _add_text(slide, spec["subtitle"], 0.82, 3.55, 7.3, 0.75, size=17, color="CBD5E1")
            if spec.get("callout"):
                _add_rect(slide, 0.82, 5.2, 3.2, 0.58, accent)
                _add_text(slide, spec["callout"], 0.96, 5.28, 2.92, 0.4, size=14, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER)
        else:
            _add_rect(slide, 0, 0, 13.333, 0.12, accent, radius=False)
            _add_text(slide, spec["title"], 0.72, 0.38, 11.9, 0.82, size=25, color="FFFFFF" if dark else "132238", bold=True)
            _add_rect(slide, 0.72, 1.25, 1.05, 0.09, accent, radius=False)

            if kind == "kpi":
                kpis = spec.get("kpis", [])
                card_width = 3.72 if len(kpis) >= 3 else 5.7
                for card_index, (value, label) in enumerate(kpis):
                    x = 0.72 + card_index * (card_width + 0.35)
                    _add_rect(slide, x, 1.72, card_width, 2.05, "FFFFFF" if not dark else "1D3557", line="DCE4EA")
                    _add_text(slide, value, x + 0.18, 2.02, card_width - 0.36, 0.75, size=27, color=accent, bold=True, align=PP_ALIGN.CENTER)
                    _add_text(slide, label, x + 0.18, 2.86, card_width - 0.36, 0.45, size=14, color="5E6C84" if not dark else "D7E3F1", align=PP_ALIGN.CENTER)
                if spec.get("bullets"):
                    for bullet_index, bullet in enumerate(spec["bullets"]):
                        _add_rect(slide, 0.76, 4.4 + bullet_index * 0.62, 0.18, 0.18, accent)
                        _add_text(slide, bullet, 1.05, 4.22 + bullet_index * 0.62, 11.4, 0.55, size=16, color="FFFFFF" if dark else "243447")

            elif kind == "chart":
                chart_spec = spec["chart"]
                chart_data = ChartData()
                chart_data.categories = chart_spec["categories"]
                for series_name, values in chart_spec["series"]:
                    chart_data.add_series(series_name, values)
                chart = slide.shapes.add_chart(
                    XL_CHART_TYPE.COLUMN_CLUSTERED,
                    Inches(0.78),
                    Inches(1.68),
                    Inches(7.4),
                    Inches(4.65),
                    chart_data,
                ).chart
                _style_chart(chart, accent)
                _add_rect(slide, 8.55, 1.68, 4.02, 4.65, "FFFFFF", line="DCE4EA")
                bullets = spec.get("bullets") or ["图表口径见输入材料"]
                for bullet_index, bullet in enumerate(bullets):
                    _add_rect(slide, 8.9, 2.05 + bullet_index * 1.35, 0.22, 0.22, accent)
                    _add_text(slide, bullet, 9.25, 1.82 + bullet_index * 1.35, 2.88, 1.08, size=15, color="243447")

            elif kind == "table":
                rows = spec["table"]
                row_count = len(rows)
                col_count = max(len(row) for row in rows)
                shape = slide.shapes.add_table(row_count, col_count, Inches(0.78), Inches(1.72), Inches(11.78), Inches(3.95))
                table = shape.table
                for row_index, row in enumerate(rows):
                    for col_index in range(col_count):
                        cell = table.cell(row_index, col_index)
                        cell.text = str(row[col_index]) if col_index < len(row) else ""
                        cell.margin_left = Inches(0.08)
                        cell.margin_right = Inches(0.08)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = _hex_color(accent if row_index == 0 else ("FFFFFF" if row_index % 2 else "EEF3F7"))
                        for paragraph in cell.text_frame.paragraphs:
                            paragraph.alignment = PP_ALIGN.CENTER if row_index == 0 else PP_ALIGN.LEFT
                            for run in paragraph.runs:
                                run.font.name = "Noto Sans CJK SC"
                                run.font.size = PptPt(12 if row_index else 13)
                                run.font.bold = row_index == 0
                                run.font.color.rgb = _hex_color("FFFFFF" if row_index == 0 else "243447")
                if spec.get("callout"):
                    _add_rect(slide, 0.78, 5.95, 11.78, 0.62, "E8F2F3", line=accent)
                    _add_text(slide, spec["callout"], 1.0, 6.04, 11.3, 0.4, size=14, color="243447", bold=True, align=PP_ALIGN.CENTER)

            elif kind == "timeline":
                items = spec.get("timeline", [])
                if items:
                    line_y = 3.75
                    _add_rect(slide, 1.15, line_y, 10.95, 0.08, accent, radius=False)
                    segment = 10.1 / max(1, len(items) - 1)
                    for item_index, (label, body) in enumerate(items):
                        x = 1.15 + item_index * segment if len(items) > 1 else 5.6
                        _add_rect(slide, x, line_y - 0.16, 0.4, 0.4, accent)
                        _add_text(slide, label, max(0.45, x - 0.78), 2.0, 1.95, 0.72, size=16, color=accent, bold=True, align=PP_ALIGN.CENTER)
                        _add_text(slide, body, max(0.35, x - 1.0), 4.3, 2.35, 1.15, size=14, color="FFFFFF" if dark else "243447", align=PP_ALIGN.CENTER)
                if spec.get("callout"):
                    _add_rect(slide, 2.8, 6.0, 7.7, 0.55, accent)
                    _add_text(slide, spec["callout"], 3.0, 6.07, 7.3, 0.36, size=13, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER)

            elif kind in {"bullets", "closing"}:
                bullets = spec.get("bullets", [])
                for bullet_index, bullet in enumerate(bullets):
                    y = 1.72 + bullet_index * 1.12
                    _add_rect(slide, 0.82, y + 0.08, 0.55, 0.55, accent)
                    _add_text(slide, f"{bullet_index + 1:02d}", 0.86, y + 0.13, 0.47, 0.35, size=12, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER)
                    _add_text(slide, bullet, 1.62, y - 0.03, 10.4, 0.8, size=18, color="FFFFFF" if dark else "243447", bold=kind == "closing")
                if spec.get("callout"):
                    _add_rect(slide, 2.2, 6.15, 8.9, 0.58, accent)
                    _add_text(slide, spec["callout"], 2.4, 6.23, 8.5, 0.38, size=14, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER)

        footer_text = deck.get("footer", "") or deck.get("author", "CoworkEvoBench")
        _add_text(slide, footer_text, 0.72, 7.08, 8.7, 0.24, size=8, color="A9B7C6" if dark else "7F8C8D")
        _add_text(slide, f"{index + 1:02d}", 11.95, 7.05, 0.6, 0.25, size=9, color=accent, bold=True, align=PP_ALIGN.RIGHT)
        if spec.get("note"):
            try:
                slide.notes_slide.notes_text_frame.text = spec["note"]
            except Exception:
                pass

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)


def _write_docx(path: Path, spec: dict[str, Any]) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    title = document.add_heading(spec["title"], 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for paragraph_text in spec.get("paragraphs", []):
        paragraph = document.add_paragraph(paragraph_text)
        paragraph.paragraph_format.space_after = Pt(7)
        for run in paragraph.runs:
            run.font.name = "Noto Sans CJK SC"
            run.font.size = Pt(10.5)
    rows = spec.get("table")
    if rows:
        table = document.add_table(rows=len(rows), cols=max(len(row) for row in rows))
        table.style = "Table Grid"
        for row_index, row in enumerate(rows):
            for col_index, value in enumerate(row):
                cell = table.cell(row_index, col_index)
                cell.text = str(value)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Noto Sans CJK SC"
                        run.font.size = Pt(9.5)
                        run.font.bold = row_index == 0
    document.core_properties.title = spec["title"]
    document.core_properties.author = "CoworkEvoBench synthetic data"
    document.save(path)


def _write_xlsx(path: Path, spec: dict[str, Any]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_spec in spec["sheets"]:
        sheet = workbook.create_sheet(sheet_spec["name"])
        for row in sheet_spec["rows"]:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="17324D")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        for column_cells in sheet.columns:
            values = [str(cell.value or "") for cell in column_cells]
            width = min(38, max(12, max(len(value) for value in values) * 1.7))
            sheet.column_dimensions[column_cells[0].column_letter].width = width
    workbook.properties.title = path.stem
    workbook.properties.creator = "CoworkEvoBench synthetic data"
    workbook.save(path)


def _write_png(path: Path, spec: dict[str, Any]) -> None:
    palette = [color.lstrip("#") for color in spec["palette"]]
    colors = [tuple(int(color[i : i + 2], 16) for i in (0, 2, 4)) for color in palette]
    width, height = 1600, 900
    image = Image.new("RGB", (width, height), colors[0])
    draw = ImageDraw.Draw(image)
    for band in range(12):
        ratio = band / 11
        color = tuple(int(colors[0][i] * (1 - ratio) + colors[1][i] * ratio) for i in range(3))
        draw.rectangle((0, band * height / 12, width, (band + 1) * height / 12), fill=color)
    draw.rounded_rectangle((170, 160, 920, 730), radius=48, fill=colors[2], outline=(255, 255, 255), width=8)
    draw.ellipse((1040, 125, 1450, 535), fill=colors[1], outline=(255, 255, 255), width=10)
    for index in range(5):
        x = 265 + index * 125
        draw.rounded_rectangle((x, 300, x + 82, 650 - index * 28), radius=18, fill=colors[1], outline=(255, 255, 255), width=4)
    image.save(path, format="PNG", optimize=True)


def _write_input(path: Path, spec: dict[str, Any]) -> None:
    kind = spec["kind"]
    if kind == "docx":
        _write_docx(path, spec)
    elif kind == "xlsx":
        _write_xlsx(path, spec)
    elif kind == "pptx":
        _render_deck(
            path,
            {
                "title": path.stem,
                "family": "synthetic-input-deck",
                "accent": spec.get("accent", "607D8B"),
                "author": "CoworkEvoBench synthetic input",
                "footer": spec.get("footer", "合成输入稿"),
                "slides": spec["slides"],
            },
        )
    elif kind == "png":
        _write_png(path, spec)
    elif kind == "json":
        path.write_text(json.dumps(spec["content"], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    elif kind == "csv":
        with path.open("w", encoding="utf-8-sig", newline="") as file:
            writer = csv.writer(file)
            writer.writerows(spec["rows"])
    elif kind == "text":
        path.write_text(spec["content"].rstrip() + "\n", encoding="utf-8")
    else:
        raise ValueError(f"不支持的输入类型：{kind}")


SOLVER_RUNTIME = r'''
import argparse
from pathlib import Path

import fitz
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


def rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def add_text(slide, text, x, y, w, h, *, size=18, color="243447", bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.06)
    frame.margin_right = Inches(0.06)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.03)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = str(text)
    run.font.name = "Noto Sans CJK SC"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return shape


def add_rect(slide, x, y, w, h, fill, *, radius=True, line=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    return shape


def style_chart(chart, accent):
    chart.has_legend = len(chart.series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.tick_labels.font.size = Pt(10)
    chart.category_axis.tick_labels.font.size = Pt(11)
    colors = [accent, "62B6CB", "F4A261", "7F8C8D"]
    for index, series in enumerate(chart.series):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = rgb(colors[index % len(colors)])


def read_source_text(input_dir: Path) -> str:
    parts = []
    for path in sorted(input_dir.iterdir()):
        suffix = path.suffix.lower()
        if suffix in {".md", ".txt", ".json", ".csv"}:
            parts.append(path.read_text(encoding="utf-8-sig"))
        elif suffix == ".docx":
            document = Document(path)
            parts.extend(paragraph.text for paragraph in document.paragraphs)
            for table in document.tables:
                parts.extend(cell.text for row in table.rows for cell in row.cells)
        elif suffix == ".xlsx":
            workbook = load_workbook(path, data_only=False, read_only=True)
            for sheet in workbook.worksheets:
                parts.extend(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None)
            workbook.close()
        elif suffix == ".pptx":
            presentation = Presentation(path)
            for slide in presentation.slides:
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False):
                        parts.append(shape.text)
                    if getattr(shape, "has_table", False):
                        parts.extend(cell.text for row in shape.table.rows for cell in row.cells)
        elif suffix == ".png":
            parts.append(path.name)
    return "\n".join(parts)


def verify_inputs(input_dir: Path) -> None:
    missing = [name for name in SPEC["input_names"] if not (input_dir / name).is_file()]
    if missing:
        raise FileNotFoundError(f"缺少输入文件：{missing}")
    source_text = read_source_text(input_dir)
    absent = [marker for marker in SPEC["source_markers"] if marker not in source_text]
    if absent:
        raise ValueError(f"输入材料与任务口径不一致，缺少：{absent}")


def build_presentation(output_path: Path, input_dir: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = SPEC["title"]
    prs.core_properties.subject = SPEC["family"]
    prs.core_properties.author = SPEC["author"]
    prs.core_properties.keywords = "CoworkEvoBench, synthetic, PowerPoint"
    accent = SPEC["accent"]
    hero_image = next(iter(sorted(input_dir.glob("*.png"))), None)

    for slide_index, item in enumerate(SPEC["slides"]):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        kind = item["kind"]
        dark = kind in {"cover", "closing"}
        bg = "132238" if dark else ("F3F6F9" if slide_index % 2 else "FFFFFF")
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(bg)

        if kind == "cover":
            add_rect(slide, 0, 0, 0.22, 7.5, accent, radius=False)
            add_rect(slide, 0.8, 0.78, 1.25, 0.13, accent, radius=False)
            if hero_image:
                slide.shapes.add_picture(str(hero_image), Inches(8.7), Inches(0.75), width=Inches(3.9), height=Inches(4.9))
                add_rect(slide, 8.7, 5.75, 3.9, 0.12, accent, radius=False)
                title_width = 7.2
            else:
                for n in range(5):
                    add_rect(slide, 9.1 + n * 0.32, 1.2 + n * 0.55, 2.5 - n * 0.18, 0.18, accent)
                title_width = 8.0
            add_text(slide, item["title"], 0.82, 1.18, title_width, 2.2, size=30, color="FFFFFF", bold=True)
            if item.get("subtitle"):
                add_text(slide, item["subtitle"], 0.82, 3.55, 7.3, 0.75, size=17, color="CBD5E1")
            if item.get("callout"):
                add_rect(slide, 0.82, 5.2, 3.2, 0.58, accent)
                add_text(slide, item["callout"], 0.96, 5.28, 2.92, 0.4, size=14, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER)
        else:
            add_rect(slide, 0, 0, 13.333, 0.12, accent, radius=False)
            add_text(slide, item["title"], 0.72, 0.38, 11.9, 0.82, size=25, color="FFFFFF" if dark else "132238", bold=True)
            add_rect(slide, 0.72, 1.25, 1.05, 0.09, accent, radius=False)

            if kind == "kpi":
                kpis = item.get("kpis", [])
                card_width = 3.72 if len(kpis) >= 3 else 5.7
                for card_index, (value, label) in enumerate(kpis):
                    x = 0.72 + card_index * (card_width + 0.35)
                    add_rect(slide, x, 1.72, card_width, 2.05, "FFFFFF" if not dark else "1D3557", line="DCE4EA")
                    add_text(slide, value, x + 0.18, 2.02, card_width - 0.36, 0.75, size=27, color=accent, bold=True, align=PP_ALIGN.CENTER)
                    add_text(slide, label, x + 0.18, 2.86, card_width - 0.36, 0.45, size=14, color="5E6C84" if not dark else "D7E3F1", align=PP_ALIGN.CENTER)
                for bullet_index, bullet in enumerate(item.get("bullets", [])):
                    add_rect(slide, 0.76, 4.4 + bullet_index * 0.62, 0.18, 0.18, accent)
                    add_text(slide, bullet, 1.05, 4.22 + bullet_index * 0.62, 11.4, 0.55, size=16, color="FFFFFF" if dark else "243447")

            elif kind == "chart":
                chart_spec = item["chart"]
                data = ChartData()
                data.categories = chart_spec["categories"]
                for series_name, values in chart_spec["series"]:
                    data.add_series(series_name, values)
                chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.78), Inches(1.68), Inches(7.4), Inches(4.65), data).chart
                style_chart(chart, accent)
                add_rect(slide, 8.55, 1.68, 4.02, 4.65, "FFFFFF", line="DCE4EA")
                bullets = item.get("bullets") or ["图表口径见输入材料"]
                for bullet_index, bullet in enumerate(bullets):
                    add_rect(slide, 8.9, 2.05 + bullet_index * 1.35, 0.22, 0.22, accent)
                    add_text(slide, bullet, 9.25, 1.82 + bullet_index * 1.35, 2.88, 1.08, size=15)

            elif kind == "table":
                rows = item["table"]
                row_count = len(rows)
                col_count = max(len(row) for row in rows)
                shape = slide.shapes.add_table(row_count, col_count, Inches(0.78), Inches(1.72), Inches(11.78), Inches(3.95))
                table = shape.table
                for row_index, row in enumerate(rows):
                    for col_index in range(col_count):
                        cell = table.cell(row_index, col_index)
                        cell.text = str(row[col_index]) if col_index < len(row) else ""
                        cell.margin_left = Inches(0.08)
                        cell.margin_right = Inches(0.08)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = rgb(accent if row_index == 0 else ("FFFFFF" if row_index % 2 else "EEF3F7"))
                        for paragraph in cell.text_frame.paragraphs:
                            paragraph.alignment = PP_ALIGN.CENTER if row_index == 0 else PP_ALIGN.LEFT
                            for run in paragraph.runs:
                                run.font.name = "Noto Sans CJK SC"
                                run.font.size = Pt(12 if row_index else 13)
                                run.font.bold = row_index == 0
                                run.font.color.rgb = rgb("FFFFFF" if row_index == 0 else "243447")
                if item.get("callout"):
                    add_rect(slide, 0.78, 5.95, 11.78, 0.62, "E8F2F3", line=accent)
                    add_text(slide, item["callout"], 1.0, 6.04, 11.3, 0.4, size=14, bold=True, align=PP_ALIGN.CENTER)

            elif kind == "timeline":
                items = item.get("timeline", [])
                if items:
                    line_y = 3.75
                    add_rect(slide, 1.15, line_y, 10.95, 0.08, accent, radius=False)
                    segment = 10.1 / max(1, len(items) - 1)
                    for item_index, (label, body) in enumerate(items):
                        x = 1.15 + item_index * segment if len(items) > 1 else 5.6
                        add_rect(slide, x, line_y - 0.16, 0.4, 0.4, accent)
                        add_text(slide, label, max(0.45, x - 0.78), 2.0, 1.95, 0.72, size=16, color=accent, bold=True, align=PP_ALIGN.CENTER)
                        add_text(slide, body, max(0.35, x - 1.0), 4.3, 2.35, 1.15, size=14, color="FFFFFF" if dark else "243447", align=PP_ALIGN.CENTER)
                if item.get("callout"):
                    add_rect(slide, 2.8, 6.0, 7.7, 0.55, accent)
                    add_text(slide, item["callout"], 3.0, 6.07, 7.3, 0.36, size=13, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER)

            elif kind in {"bullets", "closing"}:
                for bullet_index, bullet in enumerate(item.get("bullets", [])):
                    y = 1.72 + bullet_index * 1.12
                    add_rect(slide, 0.82, y + 0.08, 0.55, 0.55, accent)
                    add_text(slide, f"{bullet_index + 1:02d}", 0.86, y + 0.13, 0.47, 0.35, size=12, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER)
                    add_text(slide, bullet, 1.62, y - 0.03, 10.4, 0.8, size=18, color="FFFFFF" if dark else "243447", bold=kind == "closing")
                if item.get("callout"):
                    add_rect(slide, 2.2, 6.15, 8.9, 0.58, accent)
                    add_text(slide, item["callout"], 2.4, 6.23, 8.5, 0.38, size=14, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER)

        add_text(slide, SPEC["footer"], 0.72, 7.08, 8.7, 0.24, size=8, color="A9B7C6" if dark else "7F8C8D")
        add_text(slide, f"{slide_index + 1:02d}", 11.95, 7.05, 0.6, 0.25, size=9, color=accent, bold=True, align=PP_ALIGN.RIGHT)
        if item.get("note"):
            try:
                slide.notes_slide.notes_text_frame.text = item["note"]
            except Exception:
                pass

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)


def build_package_files(output_dir: Path) -> None:
    package_tokens = SPEC.get("package_tokens", [])
    if not package_tokens:
        return
    docx_name = next(item["name"] for item in SPEC["deliverables"] if item["type"] == "docx")
    document = Document()
    document.add_heading("2026微光志愿者大会｜讲者手册", 0)
    for heading, body in [
        ("大会主题", "让每一次善意都有回声"),
        ("时间地点", "2026年12月5日，湖滨市民中心二层大礼堂；09:00入场，09:30开始。"),
        ("讲者提示", "年度回顾用数据说明影响；故事环节不公开个人联系方式；明年计划强调志愿者保护。"),
    ]:
        document.add_heading(heading, level=1)
        document.add_paragraph(body)
    document.core_properties.title = "2026微光志愿者大会讲者手册"
    document.core_properties.author = "CoworkEvoBench synthetic solution"
    document.save(output_dir / docx_name)

    pdf_name = next(item["name"] for item in SPEC["deliverables"] if item["type"] == "pdf")
    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)
    lines = [
        "2026微光志愿者大会 参会手册",
        "主题：让每一次善意都有回声",
        "日期：2026年12月5日",
        "地点：湖滨市民中心二层大礼堂",
        "入场：09:00    开始：09:30",
        "无障碍入口：东门",
        "会务提示：请保存本手册，紧急情况联系会务台。",
    ]
    y = 90
    for index, line in enumerate(lines):
        page.insert_text((60, y), line, fontsize=20 if index == 0 else 13, fontname="china-s", color=(0.08, 0.13, 0.22))
        y += 62 if index == 0 else 38
    page.draw_rect(fitz.Rect(48, 48, 547, 794), color=(0.9, 0.49, 0.13), width=2)
    pdf.set_metadata({"title": "2026微光志愿者大会参会手册", "author": "CoworkEvoBench synthetic solution"})
    pdf.save(output_dir / pdf_name)
    pdf.close()


def main() -> int:
    parser = argparse.ArgumentParser(description="生成本题的标准交付物")
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    verify_inputs(args.input_dir)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    pptx_name = next(item["name"] for item in SPEC["deliverables"] if item["type"] == "pptx")
    build_presentation(args.output_dir / pptx_name, args.input_dir)
    build_package_files(args.output_dir)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
'''


JUDGE_RUNTIME = r'''
import argparse
import hashlib
import json
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


WEIGHTS = {"R001": 3, "R002": 2, "R003": 4, "R004": 3, "R005": 2, "R006": 2, "R007": -3, "R008": 1}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def ratio(tokens, text):
    if not tokens:
        return 1.0, []
    missing = [token for token in tokens if token not in text]
    return (len(tokens) - len(missing)) / len(tokens), missing


def inspect_presentation(path: Path):
    prs = Presentation(path)
    slide_texts = []
    charts = tables = pictures = shapes = accent_shapes = out_of_bounds = 0
    font_sizes = []
    for slide in prs.slides:
        texts = []
        for shape in slide.shapes:
            shapes += 1
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width or shape.top + shape.height > prs.slide_height:
                out_of_bounds += 1
            if getattr(shape, "has_chart", False):
                charts += 1
            if getattr(shape, "has_table", False):
                tables += 1
                texts.extend(cell.text for row in shape.table.rows for cell in row.cells)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
            if getattr(shape, "has_text_frame", False):
                texts.append(shape.text)
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.text.strip() and run.font.size:
                            font_sizes.append(run.font.size.pt)
            try:
                if shape.fill.type and str(shape.fill.fore_color.rgb).upper() == SPEC["accent"].upper():
                    accent_shapes += 1
            except Exception:
                pass
        slide_texts.append("\n".join(texts))
    return {
        "presentation": prs,
        "slide_texts": slide_texts,
        "text": "\n".join(slide_texts),
        "charts": charts,
        "tables": tables,
        "pictures": pictures,
        "shapes": shapes,
        "accent_shapes": accent_shapes,
        "out_of_bounds": out_of_bounds,
        "min_font": min(font_sizes) if font_sizes else 0,
    }


def criterion(criterion_id, score, evidence):
    return {"criterion_id": criterion_id, "score": round(max(0.0, min(1.0, score)), 4), "evidence": evidence, "method": "deterministic-pptx-v1"}


def evaluate(artifact_dir: Path):
    deliverable_paths = [artifact_dir / item["name"] for item in SPEC["deliverables"]]
    present_paths = [path for path in deliverable_paths if path.is_file()]
    hashes = [sha256(path) for path in present_paths]
    if not hashes:
        hashes = [hashlib.sha256(f"missing:{SPEC['task_id']}".encode()).hexdigest()]

    missing_files = [path.name for path in deliverable_paths if not path.is_file()]
    ppt_item = next(item for item in SPEC["deliverables"] if item["type"] == "pptx")
    ppt_path = artifact_dir / ppt_item["name"]
    inspection = None
    open_errors = []
    if ppt_path.is_file():
        try:
            inspection = inspect_presentation(ppt_path)
        except Exception as error:
            open_errors.append(f"PPTX无法打开: {error}")

    for item in SPEC["deliverables"]:
        path = artifact_dir / item["name"]
        if not path.is_file():
            continue
        try:
            if item["type"] == "docx":
                document = Document(path)
                doc_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
                absent = [token for token in SPEC.get("package_tokens", []) if token not in doc_text]
                if absent:
                    open_errors.append(f"DOCX缺少必需内容: {absent}")
            elif item["type"] == "pdf":
                if path.stat().st_size < 800 or path.read_bytes()[:5] != b"%PDF-":
                    open_errors.append("PDF文件不完整")
        except Exception as error:
            open_errors.append(f"{item['type'].upper()}无法打开: {error}")

    file_score = 1.0 if not missing_files and not open_errors and inspection else 0.0
    criteria = [criterion("R001", file_score, f"缺失={missing_files or '无'}；打开错误={open_errors or '无'}")]
    if inspection is None:
        for criterion_id in ["R002", "R003", "R004", "R005", "R006"]:
            criteria.append(criterion(criterion_id, 0.0, "无可解析的PPTX产物"))
        criteria.append(criterion("R007", 1.0, "缺少有效PPTX视为严重负向违规"))
        criteria.append(criterion("R008", 0.0, "无法校验格式保留项"))
    else:
        slide_count = len(inspection["slide_texts"])
        expected_titles = [item["title"] for item in SPEC["slides"]]
        title_hits = sum(
            index < slide_count and title in inspection["slide_texts"][index]
            for index, title in enumerate(expected_titles)
        )
        title_score = (title_hits / len(expected_titles)) if slide_count == len(expected_titles) else min(0.5, title_hits / len(expected_titles))
        criteria.append(criterion("R002", title_score, f"页数={slide_count}/{len(expected_titles)}；按顺序命中标题={title_hits}/{len(expected_titles)}"))

        content_score, missing_content = ratio(SPEC["content_tokens"], inspection["text"])
        criteria.append(criterion("R003", content_score, f"缺少内容={missing_content or '无'}"))
        data_score, missing_numbers = ratio(SPEC["numeric_tokens"], inspection["text"])
        criteria.append(criterion("R004", data_score, f"缺少数据={missing_numbers or '无'}"))

        mins = SPEC["object_mins"]
        ratios = []
        for key in ["charts", "tables", "pictures", "shapes"]:
            required = mins.get(key, 0)
            ratios.append(1.0 if required == 0 else min(1.0, inspection[key] / required))
        structure_score = sum(ratios) / len(ratios)
        criteria.append(criterion("R005", structure_score, f"图表={inspection['charts']}，表格={inspection['tables']}，图片={inspection['pictures']}，形状={inspection['shapes']}"))

        visual_checks = [inspection["out_of_bounds"] == 0, inspection["min_font"] >= 8, inspection["accent_shapes"] >= len(SPEC["slides"])]
        visual_score = sum(visual_checks) / len(visual_checks)
        criteria.append(criterion("R006", visual_score, f"越界对象={inspection['out_of_bounds']}，最小字号={inspection['min_font']:.1f}pt，强调色对象={inspection['accent_shapes']}"))

        forbidden_hits = [token for token in SPEC["forbidden_tokens"] if token in inspection["text"]]
        blank_slides = [index + 1 for index, text in enumerate(inspection["slide_texts"]) if len(text.strip()) < 12]
        violation_score = min(1.0, (len(forbidden_hits) + len(blank_slides)) / 2)
        criteria.append(criterion("R007", violation_score, f"禁止内容={forbidden_hits or '无'}；空白页={blank_slides or '无'}"))

        preserve_score, missing_preserve = ratio(SPEC["preserve_tokens"], inspection["text"])
        wide = abs(inspection["presentation"].slide_width / inspection["presentation"].slide_height - 16 / 9) < 0.02
        format_score = (preserve_score + float(wide)) / 2
        criteria.append(criterion("R008", format_score, f"缺少保留项={missing_preserve or '无'}；16:9={wide}"))

    positive_weight = sum(weight for weight in WEIGHTS.values() if weight > 0)
    scores = {item["criterion_id"]: item["score"] for item in criteria}
    earned = sum(weight * scores[criterion_id] for criterion_id, weight in WEIGHTS.items() if weight > 0)
    penalty = sum(abs(weight) * scores[criterion_id] for criterion_id, weight in WEIGHTS.items() if weight < 0)
    reward = max(0.0, min(1.0, (earned - penalty) / positive_weight))
    hurdle_failures = [] if scores["R001"] == 1.0 else ["R001"]
    if hurdle_failures:
        reward = 0.0
    result = {
        "task_id": SPEC["task_id"],
        "rubric_version": "cowork-evo-rubric-v0.1",
        "judge_version": "ppt-deterministic-v1.0",
        "artifact_sha256": sorted(set(hashes)),
        "criterion_results": criteria,
        "reward": round(reward, 4),
        "passed": not hurdle_failures and reward >= 0.7,
        "hurdle_failures": hurdle_failures,
    }
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description="对PPT任务产物做确定性评分")
    parser.add_argument("--artifact-dir", type=Path, required=True)
    parser.add_argument("--result", type=Path, required=True)
    args = parser.parse_args()
    result = evaluate(args.artifact_dir)
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"task_id": result["task_id"], "reward": result["reward"], "passed": result["passed"]}, ensure_ascii=False))
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
'''


def _task_spec(task: dict[str, Any]) -> dict[str, Any]:
    footer = task["preserve_tokens"][0] if "｜" in task["preserve_tokens"][0] else task["author"]
    if task["id"] == "ppt-005-renewal-review-revision":
        footer = "柏舟云服务｜内部资料"
    elif task["id"] == "ppt-006-warehouse-safety-layout":
        footer = "速达仓配｜安全培训"
    elif task["id"] == "ppt-007-quarterly-ops-polish":
        footer = "柚子生活｜内部经营会"
    elif task["id"] == "ppt-008-product-lines-consolidation":
        footer = "云纳智造｜内部决策资料"
    elif task["id"] == "ppt-009-architecture-executive-adaptation":
        footer = "曙光交易平台｜管理层决策"
    elif task["id"] == "ppt-010-apac-launch-localization":
        footer = "Aster Cloud｜机密"
    elif task["id"] == "ppt-011-volunteer-conference-package":
        footer = "微光公益联盟｜让善意被看见"
    elif task["id"] == "ppt-012-board-deck-audit":
        footer = "屿光能源｜董事会资料｜已核验"
    return {
        "task_id": f"cowork-evo/{task['id']}",
        "title": task["title"],
        "family": task["family"],
        "author": task["author"],
        "accent": task["accent"],
        "footer": footer,
        "input_names": [item["name"] for item in task["inputs"]],
        "source_markers": task["source_markers"],
        "deliverables": task["deliverables"],
        "slides": task["slides"],
        "content_tokens": task["content_tokens"],
        "numeric_tokens": task["numeric_tokens"],
        "preserve_tokens": task["preserve_tokens"],
        "forbidden_tokens": task["forbidden_tokens"],
        "object_mins": task["object_mins"],
        "package_tokens": task.get("package_tokens", []),
    }


def _instruction(task: dict[str, Any]) -> str:
    input_lines = "\n".join(f"- `{item['name']}`" for item in task["inputs"])
    deliverable_lines = "\n".join(f"- `{item['name']}`（{item['type'].upper()}）" for item in task["deliverables"])
    requirement_lines = "\n".join(f"- {requirement}" for requirement in task["requirements"])
    return f"""# {task['title']}

## 业务背景

你是{task['sector']}行业的{task['occupation']}。请基于提供的文件完成一份可在真实办公场景中直接交付的演示成果。本题禁止联网，所有可使用信息都在输入文件中。

## 输入文件

{input_lines}

## 任务要求

{requirement_lines}

- 每页必须有可单独理解的标题，结论、证据与行动不得互相矛盾。
- 交付物中的文字、表格、图表和图片必须可编辑、可打开且无遮挡溢出。
- 不得把 `solution/`、`tests/` 或 Rubric 内容复制进交付物。

## 交付文件

将以下文件保存到 `/app/output/`：

{deliverable_lines}

不要更改文件名，不要另行打包 ZIP。
"""


def _task_toml(task: dict[str, Any]) -> str:
    deliverable_types = ", ".join(json.dumps(item["type"]) for item in task["deliverables"])
    output_filenames = ", ".join(json.dumps(item["name"], ensure_ascii=False) for item in task["deliverables"])
    capabilities = ", ".join(json.dumps(item) for item in task["capabilities"])
    supplier_capabilities = ", ".join(
        json.dumps(item) for item in SUPPLIER_CAPABILITIES_BY_FAMILY[task["family"]]
    )
    operation_types = ", ".join(
        json.dumps(item) for item in OPERATION_TYPES_BY_FAMILY[task["family"]]
    )
    input_modalities = ", ".join(
        json.dumps(item) for item in _input_modalities(task)
    )
    keywords = ", ".join(json.dumps(item, ensure_ascii=False) for item in ["office", "ppt", task["family"], task["sector"]])
    return f'''schema_version = "1.4"
artifacts = ["/app/output", "/logs/artifacts/output"]

[task]
name = "cowork-evo/{task['id']}"
version = "0.1.0"
description = "{task['title']}，公开合成端到端PPT任务。"
keywords = [{keywords}]

[metadata]
domain = "ppt"
task_family = "{task['family']}"
taxonomy_version = "v0.1"
operation_types = [{operation_types}]
input_modalities = [{input_modalities}]
dependency_scope = "{_dependency_scope(task)}"
constraint_profiles = ["content_fidelity", "numeric_accuracy", "layout_fidelity", "style_consistency", "professional_readiness"]
sector = "{task['sector']}"
occupation = "{task['occupation']}"
deliverable_types = [{deliverable_types}]
output_filenames = [{output_filenames}]
capability_ids = [{capabilities}]
supplier_capability_ids = [{supplier_capabilities}]
difficulty_tier = "{task['difficulty']}"
multimodal = true
web_search_required = false
language = "zh-CN"
judge_types = ["deterministic"]
estimated_minutes = {task['estimated']}
provenance_id = "synthetic/{task['id']}"
leakage_group_id = "ppt-family/{task['family']}/synthetic-v1"
publishability = "public"

[agent]
timeout_sec = 1800.0

[environment]
build_timeout_sec = 1800.0
cpus = 2
memory_mb = 4096
storage_mb = 12288
gpus = 0
allow_internet = false
mcp_servers = []
workdir = "/app"

[verifier]
timeout_sec = 1800.0

[verifier.env]
JUDGE_MODEL = "deterministic-pptx-v1"
'''


def _input_modalities(task: dict[str, Any]) -> list[str]:
    modalities = {
        INPUT_MODALITY_BY_KIND[item["kind"]]
        for item in task["inputs"]
        if item["kind"] in INPUT_MODALITY_BY_KIND
    }
    if len(task["inputs"]) > 1:
        modalities.add("multi_file")
    return sorted(modalities)


def _dependency_scope(task: dict[str, Any]) -> str:
    return "multi_file" if len(task["inputs"]) > 1 else "single_file"


def _task_metadata(task: dict[str, Any]) -> dict[str, Any]:
    return {
        "task_id": f"cowork-evo/{task['id']}",
        "domain": "ppt",
        "task_family": task["family"],
        "taxonomy_version": "v0.1",
        "operation_types": OPERATION_TYPES_BY_FAMILY[task["family"]],
        "input_modalities": _input_modalities(task),
        "dependency_scope": _dependency_scope(task),
        "constraint_profiles": [
            "content_fidelity",
            "numeric_accuracy",
            "layout_fidelity",
            "style_consistency",
            "professional_readiness",
        ],
        "deliverable_types": [item["type"] for item in task["deliverables"]],
        "output_filenames": [item["name"] for item in task["deliverables"]],
        "capability_ids": task["capabilities"],
        "supplier_capability_ids": SUPPLIER_CAPABILITIES_BY_FAMILY[task["family"]],
        "sector": task["sector"],
        "occupation": task["occupation"],
        "difficulty_tier": task["difficulty"],
        "multimodal": True,
        "web_search_required": False,
        "language": "zh-CN",
        "judge_types": ["deterministic"],
        "estimated_minutes": task["estimated"],
        "provenance_id": f"synthetic/{task['id']}",
        "leakage_group_id": f"ppt-family/{task['family']}/synthetic-v1",
        "publishability": "public",
        "notes": "公开合成端到端样例；正式发布前仍需人工复核和视觉 Judge 校准。",
    }


def _rubric(task: dict[str, Any]) -> dict[str, Any]:
    files = "、".join(item["name"] for item in task["deliverables"])
    titles = "、".join(item["title"] for item in task["slides"])
    return {
        "task_id": f"cowork-evo/{task['id']}",
        "rubric_version": "cowork-evo-rubric-v0.1",
        "pass_threshold": 0.7,
        "criteria": [
            {"id": "R001", "description": f"所有必需交付文件存在、可打开且格式正确：{files}。", "weight": 3, "type": "hurdle", "dimension": "file_integrity", "method": "deterministic", "method_params": {"deliverables": task["deliverables"]}, "score_cap_on_failure": 0},
            {"id": "R002", "description": f"页数与顺序严格符合要求，且每页包含对应结论标题：{titles}。", "weight": 2, "type": "positive", "dimension": "structure", "method": "deterministic", "method_params": {"expected_slide_count": len(task["slides"])}},
            {"id": "R003", "description": "任务要求的业务内容、关键结论与行动项在演示稿中完整出现。", "weight": 4, "type": "positive", "dimension": "instruction_following", "method": "deterministic", "method_params": {"required_tokens": task["content_tokens"]}},
            {"id": "R004", "description": "交付物中的关键日期、金额、比例和统计值与输入材料一致。", "weight": 3, "type": "positive", "dimension": "data_accuracy", "method": "deterministic", "method_params": {"required_numeric_tokens": task["numeric_tokens"]}},
            {"id": "R005", "description": "图表、表格、图片和可编辑形状等演示对象达到本题的最小结构要求。", "weight": 2, "type": "positive", "dimension": "content_quality", "method": "deterministic", "method_params": task["object_mins"]},
            {"id": "R006", "description": "所有元素都在页面边界内，正文字号可读，并在全稿稳定使用指定强调色。", "weight": 2, "type": "positive", "dimension": "presentation", "method": "deterministic", "method_params": {"accent_hex": task["accent"], "min_font_pt": 8}},
            {"id": "R007", "description": "负向惩罚：交付物不得出现被明确禁止的旧值、虚假承诺、隐私信息或明显空白页。", "weight": -3, "type": "penalty", "dimension": "factuality", "method": "deterministic", "method_params": {"forbidden_tokens": task["forbidden_tokens"]}},
            {"id": "R008", "description": "演示稿保持16:9比例，并保留任务指定的品牌、页脚、期间或保密标记。", "weight": 1, "type": "positive", "dimension": "professional_standard", "method": "deterministic", "method_params": {"preserve_tokens": task["preserve_tokens"]}},
        ],
    }


def _dockerfile() -> str:
    return """FROM python:3.11-slim-bookworm

WORKDIR /app
RUN pip install --no-cache-dir \\
    python-pptx==1.0.2 \\
    openpyxl==3.1.5 \\
    python-docx==1.2.0 \\
    Pillow==11.3.0 \\
    PyMuPDF==1.26.4
RUN mkdir -p /app/output /app/data/input_files
COPY data/input_files/ /app/data/input_files/
"""


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_executable(path: Path, content: str) -> None:
    path.write_text(content, encoding="utf-8")
    mode = path.stat().st_mode
    path.chmod(mode | stat.S_IXUSR | stat.S_IXGRP | stat.S_IXOTH)


def _provenance(task: dict[str, Any], input_dir: Path) -> dict[str, Any]:
    input_hashes = {path.name: _sha256(path) for path in sorted(input_dir.iterdir()) if path.is_file()}
    reference_notes = []
    for reference in task["references"]:
        repo_name = reference.split("/", 1)[0]
        reference_notes.append(f"{reference} ({REFERENCE_URLS.get(repo_name, '本地参考仓库')})")
    hash_notes = "; ".join(f"{name}={digest}" for name, digest in input_hashes.items())
    return {
        "provenance_id": f"synthetic/{task['id']}",
        "source_type": "synthetic",
        "source_uri": f"synthetic://cowork-evolution-benchmark/{task['id']}",
        "acquired_at": ACQUIRED_AT,
        "license": "CC-BY-4.0（本项目新编合成材料）",
        "privacy_status": "通过",
        "human_rewritten": True,
        "synthetic": True,
        "publishability": "public",
        "source_sha256": list(input_hashes.values()),
        "notes": "本题所有业务名称、数据、图像和办公文件均由项目制题者合成，不含真实个人或客户数据。"
        + " 参考仓库仅用于任务形式与Rubric设计，未复制其任务文件："
        + "；".join(reference_notes)
        + f"。输入文件SHA256：{hash_notes}。",
    }


def _build_task(task: dict[str, Any]) -> None:
    root = TASKS_ROOT / task["id"]
    for directory in [
        root / "data" / "input_files",
        root / "environment",
        root / "solution",
        root / "tests",
        root / "metadata",
    ]:
        directory.mkdir(parents=True, exist_ok=True)

    input_dir = root / "data" / "input_files"
    for item in task["inputs"]:
        _write_input(input_dir / item["name"], item)

    spec = _task_spec(task)
    spec_text = pprint.pformat(spec, sort_dicts=False, width=120)
    (root / "instruction.md").write_text(_instruction(task), encoding="utf-8")
    (root / "task.toml").write_text(_task_toml(task), encoding="utf-8")
    (root / "rubric.json").write_text(json.dumps(_rubric(task), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (root / "environment" / "Dockerfile").write_text(_dockerfile(), encoding="utf-8")
    (root / "solution" / "solve.py").write_text("#!/usr/bin/env python3\n\nSPEC = " + spec_text + "\n" + SOLVER_RUNTIME, encoding="utf-8")
    _write_executable(
        root / "solution" / "solve.sh",
        """#!/usr/bin/env bash
set -euo pipefail

TASK_ROOT="$(cd "$(dirname "${BASH_SOURCE[0]}")/.." && pwd)"
INPUT_DIR="${INPUT_DIR:-${TASK_ROOT}/data/input_files}"
OUTPUT_DIR="${1:-${OUTPUT_DIR:-/app/output}}"
PYTHON_BIN="${PYTHON_BIN:-python3}"

exec "${PYTHON_BIN}" "${TASK_ROOT}/solution/solve.py" --input-dir "${INPUT_DIR}" --output-dir "${OUTPUT_DIR}"
""",
    )
    (root / "tests" / "judge.py").write_text("#!/usr/bin/env python3\n\nSPEC = " + spec_text + "\n" + JUDGE_RUNTIME, encoding="utf-8")
    _write_executable(
        root / "tests" / "test.sh",
        """#!/usr/bin/env bash
set -euo pipefail

TASK_ROOT="$(cd "$(dirname "${BASH_SOURCE[0]}")/.." && pwd)"
ARTIFACT_DIR="${1:-${ARTIFACT_DIR:-/app/output}}"
RESULT_PATH="${2:-${RESULT_PATH:-${ARTIFACT_DIR}/judge-result.json}}"
PYTHON_BIN="${PYTHON_BIN:-python3}"

exec "${PYTHON_BIN}" "${TASK_ROOT}/tests/judge.py" --artifact-dir "${ARTIFACT_DIR}" --result "${RESULT_PATH}"
""",
    )
    (root / "metadata" / "task.json").write_text(
        json.dumps(_task_metadata(task), ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    (root / "metadata" / "provenance.json").write_text(json.dumps(_provenance(task, input_dir), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> int:
    TASKS_ROOT.mkdir(parents=True, exist_ok=True)
    expected_ids = {task["id"] for task in TASKS}
    existing_ids = {path.name for path in TASKS_ROOT.iterdir() if path.is_dir()}
    unexpected = sorted(existing_ids - expected_ids)
    if unexpected:
        raise RuntimeError(f"拒绝覆盖非本脚本管理的任务目录：{unexpected}")
    for task in TASKS:
        _build_task(task)
        print(f"已生成：{task['id']} ({task['family']})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
