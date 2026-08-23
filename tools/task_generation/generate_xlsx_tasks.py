#!/usr/bin/env python3
"""生成 CoworkEvoBench 的 14 个 Excel 公开合成任务。

该脚本只生成 `tasks/xlsx/` 下的题包。所有业务数据均为固定种子的原创合成数据，
不复制参考 Benchmark 的数据文件。每个题包都包含真实可打开的输入、Golden、
可执行解法、细粒度 Rubric、确定性 Judge 与来源哈希。
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import shutil
import subprocess
import sys
import textwrap
from datetime import date, datetime, timedelta
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


REPO_ROOT = Path(__file__).resolve().parents[2]
TASKS_ROOT = REPO_ROOT / "tasks" / "xlsx"
PYTHON = Path(
    os.environ.get("COWORK_EVO_PYTHON", sys.executable)
)

NAVY = "17365D"
BLUE = "4472C4"
TEAL = "00A6A6"
GREEN = "70AD47"
ORANGE = "ED7D31"
RED = "C00000"
LIGHT_BLUE = "D9EAF7"
LIGHT_GREEN = "E2F0D9"
LIGHT_ORANGE = "FCE4D6"
LIGHT_GRAY = "E7E6E6"
WHITE = "FFFFFF"
THIN_GRAY = Side(style="thin", color="B7B7B7")


REFERENCE_REPOSITORIES = [
    {
        "name": "SpreadsheetBench-2",
        "url": "https://github.com/RUCKBReasoning/SpreadsheetBench-2",
        "files_reviewed": [
            "README.md",
            "evaluation/evaluation.py",
            "SWE-agent/config/spreadsheet.yaml",
            "SWE-agent/config/visualisation.yaml",
        ],
        "license_status": "仓库根目录未声明数据许可；仅参考题族和单元格对比方法，未复制数据",
        "release_status": "2026-08-24 时为公开 GitHub 仓库",
    },
    {
        "name": "SpreadsheetBench",
        "url": "https://github.com/EnvCommons/SpreadsheetBench",
        "files_reviewed": ["README.md", "evaluate.py"],
        "license_status": "README 标注 CC-BY-SA-4.0；仅参考 OJ 式确定评测",
        "release_status": "2026-08-24 时为公开 GitHub 仓库",
    },
    {
        "name": "OmegaUse-OfficeVal",
        "url": "https://github.com/baidu-frontier-research/OmegaUse-OfficeVal",
        "files_reviewed": [
            "README_zh-CN.md",
            "data/rubrics-en/officeval_092.json",
            "data/rubrics-en/officeval_100.json",
        ],
        "license_status": "Apache-2.0；仅参考细粒度正负 Rubric 和 Office 文件预检",
        "release_status": "2026-08-24 时为公开 GitHub 仓库",
    },
    {
        "name": "Workspace-Bench",
        "url": "https://github.com/OpenDataBox/Workspace-Bench",
        "files_reviewed": ["README.md", "docs/evaluation.md"],
        "license_status": "仓库含 LICENSE；仅参考任务隔离、细粒度 Rubric 和文件依赖设计",
        "release_status": "2026-08-24 时为公开 GitHub 仓库",
    },
    {
        "name": "Cowork-Bench",
        "url": "https://github.com/0717376/cowork_bench",
        "files_reviewed": [
            "README.md",
            "tasks/finalpool/canvas-quiz-performance-excel/docs/task.md",
            "tasks/finalpool/canvas-quiz-performance-excel/evaluation/main.py",
        ],
        "license_status": "Apache-2.0；仅参考端到端办公场景和确定性验收",
        "release_status": "2026-08-24 时为公开 GitHub 仓库",
    },
]


TASK_SPECS = [
    {
        "family": "document_to_table",
        "task_id": "xlsx-document-to-table-procurement-quotes",
        "title": "从供应商报价纪要生成采购比价表",
        "sector": "企业采购",
        "occupation": "采购专员",
        "difficulty": "intermediate",
        "minutes": 45,
        "outputs": ["采购报价结构化.xlsx"],
        "header_rows": {"报价明细": [1], "供应商汇总": [1], "最低价推荐": [1]},
        "instruction": """
        你是行政采购组的采购专员。`data/input_files/供应商报价纪要.docx`
        是询价会议记录，包含多家供应商、物料、数量、含税单价和交付周期。

        请将纪要中的六条报价整理为 `采购报价结构化.xlsx`，保存到 `/app/output`。
        工作簿必须有 `报价明细`、`供应商汇总`、`最低价推荐` 三张表。明细表按物料编码、
        含税总额升序排列，含税总额=数量×含税单价；供应商汇总表统计报价条数和金额；
        推荐表为每个物料选出含税总额最低的报价，并给出推荐合计。

        金额使用人民币格式保留两位小数；所有表启用筛选、冻结表头、统一中文表头样式，
        且不得编造纪要中没有的报价。
        """,
    },
    {
        "family": "data_cleaning",
        "task_id": "xlsx-data-cleaning-crm-leads",
        "title": "清洗并去重 CRM 客户线索",
        "sector": "企业服务",
        "occupation": "CRM 运营专员",
        "difficulty": "intermediate",
        "minutes": 50,
        "outputs": ["客户线索清洗结果.xlsx"],
        "header_rows": {"清洗结果": [1], "质量汇总": [1], "省份统计": [1]},
        "instruction": """
        你是 CRM 数据运营专员。请清洗 `客户线索_待清洗.csv`，并使用 `省份映射.xlsx`
        将省份别名统一为标准名称。以 lead_id 为主键，重复记录保留 updated_at 最晚的一条。

        输出 `客户线索清洗结果.xlsx`，包含 `清洗结果`、`质量汇总`、`省份统计` 三张表。
        手机号只保留数字，邮箱转小写并去除空格，预算统一为元，同意营销统一为“是/否”；
        对手机号、邮箱、省份和预算进行校验，在 `校验状态` 中写明“通过”或具体问题。
        不要删除校验不通过的唯一线索；金额保留两位小数，表头冻结且可筛选。
        """,
    },
    {
        "family": "table_transformation",
        "task_id": "xlsx-table-transformation-regional-sales",
        "title": "将区域月度宽表转为可分析长表",
        "sector": "零售",
        "occupation": "销售分析师",
        "difficulty": "intermediate",
        "minutes": 45,
        "outputs": ["区域月度销售长表.xlsx"],
        "header_rows": {"销售长表": [1], "区域汇总": [1], "产品汇总": [1]},
        "instruction": """
        `区域月度销售_宽表.xlsx` 把 2026 年 1–3 月的销售额和销量放在不同列中。
        请输出 `区域月度销售长表.xlsx`，把它转为每行一个“区域-产品-月份”观测的长表。

        `销售长表` 必须包含区域、产品、月份、销售额、销量、客单价六列，客单价=销售额/销量。
        再生成 `区域汇总` 和 `产品汇总`，分别统计销售额、销量和加权客单价。
        保留数据精度，不得对月份或类别进行丢失性合并；三张表均需标准化表头、筛选和冻结窗格。
        """,
    },
    {
        "family": "formula_authoring",
        "task_id": "xlsx-formula-authoring-order-margin",
        "title": "为订单明细补齐销售额和毛利公式",
        "sector": "电商",
        "occupation": "商业分析师",
        "difficulty": "intermediate",
        "minutes": 40,
        "outputs": ["订单毛利计算.xlsx"],
        "header_rows": {"订单明细": [1], "经营摘要": [1]},
        "instruction": """
        `订单明细_待计算.xlsx` 的基础字段已完整，H–K 列和 `经营摘要` 中的计算区为空。
        请输出 `订单毛利计算.xlsx`，为每行写入可编辑 Excel 公式：
        销售额=数量×含税单价×(1-折扣率)，成本=数量×单位成本，毛利=销售额-成本，毛利率=毛利/销售额。

        `经营摘要` 必须用公式计算订单数、总销售额、总成本、总毛利和整体毛利率。
        不得把计算结果直接写成硬编码数值，不得改动 A–G 列的订单基础数据；
        金额、百分比格式要正确，且不能出现 `#REF!` 或 `#DIV/0!`。
        """,
    },
    {
        "family": "model_building",
        "task_id": "xlsx-model-building-saas-plan",
        "title": "搭建 SaaS 三年经营模型",
        "sector": "SaaS",
        "occupation": "FP&A 分析师",
        "difficulty": "advanced",
        "minutes": 90,
        "outputs": ["SaaS三年经营模型.xlsx"],
        "header_rows": {"三年模型": [2], "指标摘要": [1]},
        "instruction": """
        你是 SaaS 公司的 FP&A 分析师。请使用 `SaaS经营假设与历史.xlsx` 中的 2025 年实际数和
        2026–2028 年预测假设，输出 `SaaS三年经营模型.xlsx`。

        `三年模型` 需用跨表公式计算期初客户、新增客户、流失客户、期末客户、平均客户、
        订阅收入、销售成本、毛利、薪酬、营销费用、其他费用、EBITDA 和 EBITDA 率。
        流失客户=期初客户×年流失率，收入=平均客户×月 ARPU×12。
        `指标摘要` 展示收入 CAGR、2028 年毛利率和 EBITDA 率。
        假设单元格使用浅蓝色，公式单元格使用白色，百分比与金额显示格式正确。
        """,
    },
    {
        "family": "scenario_analysis",
        "task_id": "xlsx-scenario-analysis-logistics-pricing",
        "title": "分析物流价格与业务量情景",
        "sector": "物流",
        "occupation": "定价分析师",
        "difficulty": "advanced",
        "minutes": 75,
        "outputs": ["物流定价情景分析.xlsx"],
        "header_rows": {"情景分析": [1], "敏感性矩阵": [2]},
        "instruction": """
        `物流定价假设.xlsx` 包含基准包裹量、客单价、单票变动成本、固定成本和价格弹性，
        以及悲观、基准、乐观三组调价幅度。请输出 `物流定价情景分析.xlsx`。

        `情景分析` 需用公式给出各情景的调整后价格、预计包裹量、收入、变动成本、贡献毛利和营业利润。
        业务量=基准包裹量×(1+价格弹性×调价幅度)。
        `敏感性矩阵` 必须横向展示 -10%、-5%、0%、5%、10% 的调价，纵向展示 -10%、0%、10% 的额外业务量变动，
        每个交叉单元格计算营业利润。正负利润要有可读的条件格式。
        """,
    },
    {
        "family": "pivot_analysis",
        "task_id": "xlsx-pivot-analysis-retail-channel",
        "title": "制作零售渠道透视式分析",
        "sector": "零售",
        "occupation": "渠道分析师",
        "difficulty": "intermediate",
        "minutes": 60,
        "outputs": ["零售渠道透视分析.xlsx"],
        "header_rows": {"交易明细": [1], "透视分析": [2], "月度趋势": [1]},
        "instruction": """
        请分析 `零售交易明细.xlsx`，输出 `零售渠道透视分析.xlsx`。
        由于跨平台的原生 PivotTable 对象兼容性不稳定，本题要求生成可编辑的“数据透视式汇总”，
        无需原生 PivotTable 缓存，但必须从明细记录真实聚合，不得只贴图片。

        工作簿包含 `交易明细`、`透视分析`、`月度趋势`。`透视分析` 以渠道为行、品类为列，
        展示销售额，附行合计、列合计和各渠道毛利率。`月度趋势` 按月汇总销售额、毛利和毛利率，
        并添加一张可编辑的月度销售额柱状图。
        """,
    },
    {
        "family": "charting",
        "task_id": "xlsx-charting-growth-funnel",
        "title": "为增长漏斗周报制作可编辑图表",
        "sector": "互联网服务",
        "occupation": "增长分析师",
        "difficulty": "intermediate",
        "minutes": 45,
        "outputs": ["增长漏斗图表.xlsx"],
        "header_rows": {"周度数据": [1], "图表看板": [1]},
        "instruction": """
        请使用 `增长漏斗周报.xlsx` 生成 `增长漏斗图表.xlsx`。保留一张 `周度数据` 明细表，
        新建 `图表看板`，显示累计线索、累计商机、累计下单和整体线索转订单率。

        看板至少包含两张可编辑图表：一张以周次为横轴的线索/商机/订单趋势折线图，
        一张展示“线索→商机”和“商机→订单”转化率的分组柱状图。
        图表必须有清晰标题、图例和百分比坐标格式，不得使用静态截图代替图表对象。
        """,
    },
    {
        "family": "dashboard",
        "task_id": "xlsx-dashboard-sales-operations",
        "title": "制作销售运营管理仪表盘",
        "sector": "B2B 销售",
        "occupation": "销售运营经理",
        "difficulty": "advanced",
        "minutes": 90,
        "outputs": ["销售运营仪表盘.xlsx"],
        "header_rows": {"明细": [1], "分析": [1], "仪表盘": [1]},
        "instruction": """
        你是销售运营经理。请根据 `销售订单与目标.xlsx` 输出 `销售运营仪表盘.xlsx`，
        用于月度经营会。工作簿包含 `明细`、`分析`、`仪表盘` 三张表。

        仪表盘顶部展示总销售额、总毛利、整体毛利率、订单数和目标达成率五个 KPI；
        下方至少包含月度销售趋势、区域目标达成率、产品毛利贡献三张可编辑图表。
        `分析` 表提供图表所依赖的可核验汇总；金额、百分比、配色和条件格式要统一，
        不得为了达成目标而修改输入订单。
        """,
    },
    {
        "family": "template_completion",
        "task_id": "xlsx-template-completion-headcount-plan",
        "title": "完成人员编制与招聘计划模板",
        "sector": "人力资源",
        "occupation": "HRBP",
        "difficulty": "intermediate",
        "minutes": 50,
        "outputs": ["人员编制计划完成版.xlsx"],
        "header_rows": {"岗位需求": [3], "部门汇总": [1]},
        "instruction": """
        `人员编制计划模板.xlsx` 已包含公司标题、样式、数据验证和隐藏的 `说明` 表。
        请把 `招聘需求.csv` 写入模板，输出 `人员编制计划完成版.xlsx`。

        `岗位需求` 的 A–G 列保持指定字段顺序，其中招聘缺口=目标编制-现有人数，使用公式计算；
        缺口大于 2 或到岗日早于 2026-10-01 的行标注为高优先级。新建 `部门汇总` 汇总目标编制、
        现有人数和招聘缺口。必须保留原有合并标题、列宽、颜色和隐藏的 `说明` 表。
        """,
    },
    {
        "family": "spreadsheet_repair",
        "task_id": "xlsx-spreadsheet-repair-budget-errors",
        "title": "修复年度预算工作簿中的公式和样式错误",
        "sector": "企业财务",
        "occupation": "财务分析师",
        "difficulty": "advanced",
        "minutes": 65,
        "outputs": ["年度预算修复版.xlsx"],
        "header_rows": {"部门预算": [1], "汇总": [1], "修复日志": [1]},
        "instruction": """
        `年度预算_公式损坏.xlsx` 中存在 `#REF!`、错误求和范围、错位的年度差异公式，
        以及一个表头样式损坏。请输出 `年度预算修复版.xlsx`。

        修复 `部门预算` 中每个部门的年度合计和与获批预算的差异，修复合计行；
        修复 `汇总` 对各部门的引用，不得替换为硬编码数值。新建 `修复日志`，
        记录每个修复单元格、原公式/值、修复后公式/值和原因。工作簿中不得残留任何 Excel 错误标记，
        且要恢复表头的统一深蓝色样式。
        """,
    },
    {
        "family": "spreadsheet_audit",
        "task_id": "xlsx-spreadsheet-audit-project-costs",
        "title": "审计项目成本台账并形成可追溯发现",
        "sector": "工程项目",
        "occupation": "内审专员",
        "difficulty": "advanced",
        "minutes": 80,
        "outputs": ["项目成本审计.xlsx"],
        "header_rows": {"原始台账": [1], "审计发现": [1], "审计汇总": [1]},
        "instruction": """
        你是内审专员。请对 `项目成本台账_待审计.xlsx` 与 `采购合同限额.xlsx` 进行交叉审计，
        输出 `项目成本审计.xlsx`。

        至少检查五类风险：重复发票号、单笔金额超合同单笔限额、累计金额超合同总额、
        缺失审批人、负数金额。原始数据原样复制到 `原始台账`，不得静默修改原记录。
        `审计发现` 每行包含发现编号、风险类型、风险级别、交易编号、证据和建议；
        `审计汇总` 统计各类发现数和涉及金额。高风险行用红色突出，所有发现必须能回溯到输入行。
        """,
    },
    {
        "family": "multi_workbook_integration",
        "task_id": "xlsx-multi-workbook-integration-branch-orders",
        "title": "合并三家分公司订单并关联产品主数据",
        "sector": "分销",
        "occupation": "数据分析师",
        "difficulty": "advanced",
        "minutes": 75,
        "outputs": ["分公司订单汇总.xlsx"],
        "header_rows": {"合并明细": [1], "分公司汇总": [1], "对账清单": [1]},
        "instruction": """
        `华东订单.xlsx`、`华南订单.xlsx`、`华北订单.xlsx` 中存在跨文件重复订单，
        `产品主数据.xlsx` 提供品类和标准成本。请输出 `分公司订单汇总.xlsx`。

        合并时以订单号为唯一键，如有重复则保留更新时间最晚的记录，并记录保留来源。
        `合并明细` 关联品类和标准成本，计算销售额、标准成本和标准毛利；
        `分公司汇总` 按来源分公司统计订单数、销售额和毛利；` 对账清单` 列出被去重的记录及原因。
        任一未在主数据中的 SKU 必须显式标记，不得臆测成本。
        """,
    },
    {
        "family": "analysis_to_report",
        "task_id": "xlsx-analysis-to-report-service-review",
        "title": "将客户服务数据分析转化为经营报告",
        "sector": "客户服务",
        "occupation": "服务运营分析师",
        "difficulty": "advanced",
        "minutes": 90,
        "outputs": ["客户服务经营分析.xlsx", "客户服务经营分析报告.docx"],
        "header_rows": {"月度数据": [1], "KPI摘要": [1], "趋势分析": [1]},
        "instruction": """
        你是客户服务运营分析师。请使用 `客户服务月度数据.xlsx` 生成两个交付物：
        `客户服务经营分析.xlsx` 和 `客户服务经营分析报告.docx`，均保存到 `/app/output`。

        Excel 包含 `月度数据`、`KPI摘要`、`趋势分析`，计算全年工单量、平均 SLA 达成率、
        加权平均 CSAT、流失客户数、收入、服务成本、利润和利润率，并用可编辑图表显示 SLA/CSAT 与利润趋势。
        Word 报告必须包含“执行摘要”、“关键发现”、“改进建议”三章，引用 Excel 中一致的全年 KPI，
        指出 SLA 最低月份和利润最低月份，并提出至少三条有针对性的建议。不得引入输入数据无法支持的市场事实。
        """,
    },
]


def clean_text(value: str) -> str:
    return textwrap.dedent(value).strip() + "\n"


def write_text(path: Path, value: str, executable: bool = False) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(value, encoding="utf-8")
    if executable:
        path.chmod(0o755)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def style_header(ws, row: int = 1, start_col: int = 1, end_col: int | None = None) -> None:
    end_col = end_col or ws.max_column
    for cell in ws.iter_cols(min_col=start_col, max_col=end_col, min_row=row, max_row=row):
        item = cell[0]
        item.font = Font(name="微软雅黑", size=10, bold=True, color=WHITE)
        item.fill = PatternFill("solid", fgColor=NAVY)
        item.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        item.border = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)
    ws.row_dimensions[row].height = 28


def style_body(ws, min_row: int = 2, max_row: int | None = None) -> None:
    max_row = max_row or ws.max_row
    for row in ws.iter_rows(min_row=min_row, max_row=max_row):
        for cell in row:
            cell.font = Font(name="微软雅黑", size=10, color="222222")
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.border = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)


def finish_sheet(ws, header_row: int = 1, freeze: str | None = None) -> None:
    style_header(ws, header_row)
    style_body(ws, header_row + 1)
    ws.freeze_panes = freeze or f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
    for column in range(1, ws.max_column + 1):
        letter = get_column_letter(column)
        values = [str(ws.cell(row, column).value or "") for row in range(1, min(ws.max_row, 80) + 1)]
        width = min(max(max((len(v) for v in values), default=8) * 1.2 + 2, 10), 30)
        ws.column_dimensions[letter].width = width
    ws.sheet_view.showGridLines = False


def add_excel_table(ws, name: str, ref: str | None = None) -> None:
    ref = ref or f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def save_workbook(wb: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(path)


def task_root(spec: dict) -> Path:
    return TASKS_ROOT / spec["task_id"]


TASK_METADATA_PROFILES = {
    "document_to_table": {
        "operation_types": ["transform", "analyze"],
        "input_modalities": ["document"],
        "dependency_scope": "single_file",
        "constraint_profiles": ["content_fidelity", "numeric_accuracy", "traceability", "professional_readiness"],
        "capability_ids": ["K1", "K2", "K3", "K6"],
        "supplier_capability_ids": ["C2", "C7", "C9", "C10", "C12"],
    },
    "data_cleaning": {
        "operation_types": ["transform", "analyze"],
        "input_modalities": ["spreadsheet", "plain_text", "multi_file"],
        "dependency_scope": "multi_file",
        "constraint_profiles": ["content_fidelity", "numeric_accuracy", "traceability"],
        "capability_ids": ["K2", "K3", "K10", "K12"],
        "supplier_capability_ids": ["C2", "C10", "C12"],
    },
    "table_transformation": {
        "operation_types": ["transform", "analyze"],
        "input_modalities": ["spreadsheet"],
        "dependency_scope": "single_file",
        "constraint_profiles": ["content_fidelity", "numeric_accuracy", "professional_readiness"],
        "capability_ids": ["K2", "K3", "K4"],
        "supplier_capability_ids": ["C2", "C10"],
    },
    "formula_authoring": {
        "operation_types": ["edit", "create"],
        "input_modalities": ["spreadsheet"],
        "dependency_scope": "single_file",
        "constraint_profiles": ["numeric_accuracy", "formula_correctness", "preservation"],
        "capability_ids": ["K2", "K4", "K7"],
        "supplier_capability_ids": ["C8", "C10"],
    },
    "model_building": {
        "operation_types": ["create", "analyze"],
        "input_modalities": ["spreadsheet"],
        "dependency_scope": "single_file",
        "constraint_profiles": ["numeric_accuracy", "formula_correctness", "traceability", "professional_readiness"],
        "capability_ids": ["K2", "K4", "K5", "K8"],
        "supplier_capability_ids": ["C8", "C10"],
    },
    "scenario_analysis": {
        "operation_types": ["analyze", "create"],
        "input_modalities": ["spreadsheet"],
        "dependency_scope": "single_file",
        "constraint_profiles": ["numeric_accuracy", "formula_correctness", "traceability", "professional_readiness"],
        "capability_ids": ["K2", "K4", "K5", "K8"],
        "supplier_capability_ids": ["C8", "C10"],
    },
    "pivot_analysis": {
        "operation_types": ["analyze", "create"],
        "input_modalities": ["spreadsheet"],
        "dependency_scope": "single_file",
        "constraint_profiles": ["numeric_accuracy", "traceability", "professional_readiness"],
        "capability_ids": ["K2", "K3", "K5", "K8"],
        "supplier_capability_ids": ["C2", "C5", "C8", "C10"],
    },
    "charting": {
        "operation_types": ["create", "analyze"],
        "input_modalities": ["spreadsheet"],
        "dependency_scope": "single_file",
        "constraint_profiles": ["numeric_accuracy", "layout_fidelity", "style_consistency", "professional_readiness"],
        "capability_ids": ["K2", "K5", "K8", "K9"],
        "supplier_capability_ids": ["C5", "C10"],
    },
    "dashboard": {
        "operation_types": ["create", "analyze"],
        "input_modalities": ["spreadsheet"],
        "dependency_scope": "single_file",
        "constraint_profiles": ["numeric_accuracy", "formula_correctness", "layout_fidelity", "professional_readiness"],
        "capability_ids": ["K2", "K4", "K5", "K8"],
        "supplier_capability_ids": ["C2", "C5", "C8", "C10"],
    },
    "template_completion": {
        "operation_types": ["edit", "create"],
        "input_modalities": ["spreadsheet", "plain_text", "multi_file"],
        "dependency_scope": "multi_file",
        "constraint_profiles": ["formula_correctness", "layout_fidelity", "preservation", "style_consistency"],
        "capability_ids": ["K1", "K4", "K7", "K9", "K10"],
        "supplier_capability_ids": ["C2", "C8", "C10", "C12"],
    },
    "spreadsheet_repair": {
        "operation_types": ["repair", "audit", "edit"],
        "input_modalities": ["spreadsheet"],
        "dependency_scope": "single_file",
        "constraint_profiles": ["numeric_accuracy", "formula_correctness", "preservation", "traceability"],
        "capability_ids": ["K2", "K4", "K7", "K12"],
        "supplier_capability_ids": ["C8", "C10"],
    },
    "spreadsheet_audit": {
        "operation_types": ["audit", "analyze"],
        "input_modalities": ["spreadsheet", "multi_file"],
        "dependency_scope": "multi_file",
        "constraint_profiles": ["content_fidelity", "numeric_accuracy", "traceability", "professional_readiness"],
        "capability_ids": ["K2", "K5", "K10", "K12"],
        "supplier_capability_ids": ["C2", "C8", "C10"],
    },
    "multi_workbook_integration": {
        "operation_types": ["transform", "analyze"],
        "input_modalities": ["spreadsheet", "multi_file"],
        "dependency_scope": "multi_file",
        "constraint_profiles": ["content_fidelity", "numeric_accuracy", "cross_artifact_consistency", "traceability"],
        "capability_ids": ["K2", "K3", "K10", "K12"],
        "supplier_capability_ids": ["C2", "C8", "C10"],
    },
    "analysis_to_report": {
        "operation_types": ["analyze", "create"],
        "input_modalities": ["spreadsheet"],
        "dependency_scope": "cross_app",
        "constraint_profiles": ["content_fidelity", "numeric_accuracy", "cross_artifact_consistency", "professional_readiness"],
        "capability_ids": ["K1", "K2", "K5", "K6", "K8", "K13"],
        "supplier_capability_ids": ["C2", "C5", "C7", "C8", "C10", "C12"],
    },
}


def task_metadata(spec: dict) -> dict:
    profile = TASK_METADATA_PROFILES[spec["family"]]
    deliverable_types = sorted({Path(name).suffix.lstrip(".") for name in spec["outputs"]})
    return {
        "task_id": f"cowork-evo/{spec['task_id']}",
        "domain": "xlsx",
        "task_family": spec["family"],
        "taxonomy_version": "v0.1",
        "operation_types": profile["operation_types"],
        "input_modalities": profile["input_modalities"],
        "dependency_scope": profile["dependency_scope"],
        "constraint_profiles": profile["constraint_profiles"],
        "sector": spec["sector"],
        "occupation": spec["occupation"],
        "deliverable_types": deliverable_types,
        "output_filenames": spec["outputs"],
        "capability_ids": profile["capability_ids"],
        "supplier_capability_ids": profile["supplier_capability_ids"],
        "difficulty_tier": spec["difficulty"],
        "multimodal": False,
        "web_search_required": False,
        "language": "zh-CN",
        "judge_types": ["deterministic"],
        "estimated_minutes": spec["minutes"],
        "provenance_id": f"{spec['task_id']}-synthetic-v1",
        "leakage_group_id": f"public-synthetic-{spec['family']}-001",
        "publishability": "public",
        "notes": "公开合成端到端 Excel 样例；正式发布前仍需人工和 LibreOffice 视觉复核。",
    }


def make_base_task(spec: dict) -> Path:
    root = task_root(spec)
    if root.exists():
        shutil.rmtree(root)
    for relative in [
        "environment",
        "data/input_files",
        "solution",
        "tests",
        "metadata",
    ]:
        (root / relative).mkdir(parents=True, exist_ok=True)

    write_text(root / "instruction.md", clean_text(spec["instruction"]))
    keywords = ["office", "xlsx", spec["family"], "synthetic-public"]
    metadata = task_metadata(spec)
    profile = TASK_METADATA_PROFILES[spec["family"]]
    task_toml = f'''schema_version = "1.4"
artifacts = ["/app/output", "/logs/artifacts/output"]

[task]
name = "cowork-evo/{spec['task_id']}"
version = "0.1.0"
description = "{spec['title']}"
keywords = {json.dumps(keywords, ensure_ascii=False)}

[metadata]
domain = "xlsx"
task_family = "{spec['family']}"
taxonomy_version = "v0.1"
operation_types = {json.dumps(profile['operation_types'], ensure_ascii=False)}
input_modalities = {json.dumps(profile['input_modalities'], ensure_ascii=False)}
dependency_scope = "{profile['dependency_scope']}"
constraint_profiles = {json.dumps(profile['constraint_profiles'], ensure_ascii=False)}
deliverable_types = {json.dumps(metadata['deliverable_types'], ensure_ascii=False)}
output_filenames = {json.dumps(spec['outputs'], ensure_ascii=False)}
capability_ids = {json.dumps(profile['capability_ids'], ensure_ascii=False)}
supplier_capability_ids = {json.dumps(profile['supplier_capability_ids'], ensure_ascii=False)}
sector = "{spec['sector']}"
occupation = "{spec['occupation']}"
difficulty_tier = "{spec['difficulty']}"
multimodal = false
web_search_required = false
language = "zh-CN"
judge_types = ["deterministic"]
estimated_minutes = {spec['minutes']}
provenance_id = "{spec['task_id']}-synthetic-v1"
leakage_group_id = "public-synthetic-{spec['family']}-001"
publishability = "public"

[agent]
timeout_sec = 1800.0

[environment]
build_timeout_sec = 1800.0
cpus = 2
memory_mb = 4096
storage_mb = 12288
gpus = 0
allow_internet = false
mcp_servers = []
workdir = "/app"

[verifier]
timeout_sec = 600.0

[verifier.env]
PYTHONUNBUFFERED = "1"
'''
    write_text(root / "task.toml", task_toml)
    write_text(
        root / "metadata" / "task.json",
        json.dumps(metadata, ensure_ascii=False, indent=2) + "\n",
    )

    dockerfile = '''FROM python:3.12-slim

WORKDIR /app
RUN python -m pip install --no-cache-dir openpyxl==3.1.5 python-docx==1.2.0
RUN mkdir -p /app/data/input_files /app/output
COPY data/input_files/ /app/data/input_files/
'''
    write_text(root / "environment" / "Dockerfile", dockerfile)

    solve_sh = '''#!/usr/bin/env bash
set -euo pipefail

TASK_ROOT="$(cd "$(dirname "${BASH_SOURCE[0]}")/.." && pwd)"
INPUT_DIR="${INPUT_DIR:-${TASK_ROOT}/data/input_files}"
if [[ -n "${1:-}" ]]; then
  OUTPUT_DIR="$1"
elif [[ -d /app ]]; then
  OUTPUT_DIR="${OUTPUT_DIR:-/app/output}"
else
  OUTPUT_DIR="${OUTPUT_DIR:-${TASK_ROOT}/output}"
fi
mkdir -p "$OUTPUT_DIR"
python3 "${TASK_ROOT}/solution/solve.py" --input-dir "$INPUT_DIR" --output-dir "$OUTPUT_DIR"
'''
    write_text(root / "solution" / "solve.sh", solve_sh, executable=True)

    test_sh = '''#!/usr/bin/env bash
set -euo pipefail

TASK_ROOT="$(cd "$(dirname "${BASH_SOURCE[0]}")/.." && pwd)"
CANDIDATE_DIR="${1:-${CANDIDATE_DIR:-/app/output}}"
python3 "${TASK_ROOT}/tests/judge.py" \
  --candidate-dir "$CANDIDATE_DIR" \
  --golden-dir "${TASK_ROOT}/solution" \
  --checks "${TASK_ROOT}/tests/checks.json"
'''
    write_text(root / "tests" / "test.sh", test_sh, executable=True)
    return root


def rubric_for(spec: dict) -> dict:
    task_id = f"cowork-evo/{spec['task_id']}"
    output_names = "、".join(spec["outputs"])
    return {
        "task_id": task_id,
        "rubric_version": "cowork-evo-rubric-v0.1",
        "pass_threshold": 0.8,
        "criteria": [
            {
                "id": "R001",
                "description": f"必需交付文件 {output_names} 全部存在，扩展名正确且可正常打开。",
                "weight": 2,
                "type": "hurdle",
                "dimension": "file_integrity",
                "method": "deterministic",
                "method_params": {"required_files": spec["outputs"]},
                "score_cap_on_failure": 0,
            },
            {
                "id": "R002",
                "description": "工作簿工作表、表头、行列结构和必要报告章节完整，与任务指定布局一致。",
                "weight": 2,
                "type": "positive",
                "dimension": "structure",
                "method": "deterministic",
                "method_params": {},
            },
            {
                "id": "R003",
                "description": "所有关键数字、分组汇总、去重结果与输入文件一致，数值误差不超过 0.01。",
                "weight": 4,
                "type": "positive",
                "dimension": "data_accuracy",
                "method": "deterministic",
                "method_params": {"numeric_tolerance": 0.01},
            },
            {
                "id": "R004",
                "description": "指定为动态计算的单元格使用正确、可编辑的 Excel 公式，引用范围正确。",
                "weight": 3,
                "type": "positive",
                "dimension": "data_accuracy",
                "method": "deterministic",
                "method_params": {},
            },
            {
                "id": "R005",
                "description": "表头、数字格式、冻结窗格、筛选、条件格式和可编辑图表符合任务的专业呈现要求。",
                "weight": 2,
                "type": "positive",
                "dimension": "presentation",
                "method": "deterministic",
                "method_params": {},
            },
            {
                "id": "R006",
                "description": "交付物没有空白工作表、断裂的图表数据源或 #REF!/#DIV/0! 等明显错误。",
                "weight": 1,
                "type": "positive",
                "dimension": "professional_standard",
                "method": "deterministic",
                "method_params": {},
            },
            {
                "id": "R007",
                "description": "如果交付物空白、擅自改动受保护的输入基础数据，或编造输入无法支持的关键事实，触发负向惩罚。",
                "weight": -4,
                "type": "penalty",
                "dimension": "factuality",
                "method": "deterministic",
                "method_params": {},
            },
        ],
    }


COMMON_SOLVER = r'''
from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "17365D"
BLUE = "4472C4"
GREEN = "70AD47"
ORANGE = "ED7D31"
RED = "C00000"
WHITE = "FFFFFF"
LIGHT_BLUE = "D9EAF7"
THIN = Side(style="thin", color="B7B7B7")


def header(ws, row=1):
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.font = Font(name="微软雅黑", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ws.row_dimensions[row].height = 28


def body(ws, start_row=2):
    for row in ws.iter_rows(min_row=start_row):
        for cell in row:
            cell.font = Font(name="微软雅黑", size=10, color="222222")
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def finish(ws, header_row=1, freeze=None, autofilter=True):
    header(ws, header_row)
    body(ws, header_row + 1)
    ws.freeze_panes = freeze or f"A{header_row + 1}"
    if autofilter:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.sheet_view.showGridLines = False
    for col in range(1, ws.max_column + 1):
        vals = [str(ws.cell(r, col).value or "") for r in range(1, min(ws.max_row, 80) + 1)]
        width = min(max(max((len(v) for v in vals), default=8) * 1.2 + 2, 10), 30)
        ws.column_dimensions[get_column_letter(col)].width = width


def table(ws, name, ref=None):
    ref = ref or f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    obj = Table(displayName=name, ref=ref)
    obj.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(obj)


def money(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = '¥#,##0.00;[Red]-¥#,##0.00'


def percent(ws, columns, start_row=2, end_row=None):
    end_row = end_row or ws.max_row
    for col in columns:
        for row in range(start_row, end_row + 1):
            ws[f"{col}{row}"].number_format = "0.00%"


def save(wb, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(path)
'''


JUDGE_SOURCE = r'''#!/usr/bin/env python3
"""确定性 Office 任务 Judge。

Golden 和 checks 属于评测侧隐藏资产，不得挂载到 Agent 可见的输入目录。
"""
from __future__ import annotations

import argparse
import json
import math
import re
import sys
from datetime import date, datetime, time
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

ERROR_MARKERS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#NULL!", "#N/A", "#NUM!"}


def normalize_formula(value):
    if not isinstance(value, str) or not value.startswith("="):
        return value
    return re.sub(r"\s+", "", value.replace("$", "").upper())


def normalize(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def equivalent(actual, expected, tolerance):
    if isinstance(actual, str) and actual.startswith("=") and isinstance(expected, str) and expected.startswith("="):
        return normalize_formula(actual) == normalize_formula(expected)
    actual, expected = normalize(actual), normalize(expected)
    if actual in (None, "") and expected in (None, ""):
        return True
    if isinstance(actual, bool) or isinstance(expected, bool):
        return actual is expected
    if isinstance(actual, (int, float)) and isinstance(expected, (int, float)):
        if math.isnan(actual) or math.isnan(expected):
            return False
        return abs(float(actual) - float(expected)) <= tolerance
    return actual == expected


def color_rgb(color):
    if color is None:
        return None
    if getattr(color, "type", None) == "rgb" and isinstance(color.rgb, str):
        return color.rgb[-6:].upper()
    if getattr(color, "type", None) == "theme":
        return f"THEME:{color.theme}"
    return None


def compare_range(actual_ws, golden_ws, range_ref, tolerance, errors):
    for row in golden_ws[range_ref]:
        for golden_cell in row:
            actual_cell = actual_ws[golden_cell.coordinate]
            if not equivalent(actual_cell.value, golden_cell.value, tolerance):
                errors.append(
                    f"{actual_ws.title}!{golden_cell.coordinate}: "
                    f"期望 {golden_cell.value!r}，实际 {actual_cell.value!r}"
                )


def chart_type(chart):
    return type(chart).__name__


def scan_errors(wb):
    found = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if value in ERROR_MARKERS or (
                    isinstance(value, str) and value.startswith("=") and any(marker in value for marker in ERROR_MARKERS)
                ):
                    found.append(f"{ws.title}!{cell.coordinate}={value}")
    return found


def check_xlsx(candidate, golden, config, tolerance):
    result = {"integrity": [], "structure": [], "accuracy": [], "formula": [], "format": [], "quality": []}
    try:
        actual_wb = load_workbook(candidate, data_only=False)
        golden_wb = load_workbook(golden, data_only=False)
        result["integrity"].append("工作簿可正常打开")
    except Exception as exc:
        result["integrity"].append(f"ERROR:工作簿无法打开: {exc}")
        return result

    required = config.get("sheets", [])
    for name in required:
        if name not in actual_wb.sheetnames:
            result["structure"].append(f"ERROR:缺少工作表 {name}")
        else:
            result["structure"].append(f"工作表 {name} 存在")

    if config.get("exact_sheet_order") and actual_wb.sheetnames[: len(required)] != required:
        result["structure"].append(
            f"ERROR:工作表顺序不正确: {actual_wb.sheetnames}"
        )

    for item in config.get("ranges", []):
        sheet = item["sheet"]
        if sheet not in actual_wb.sheetnames or sheet not in golden_wb.sheetnames:
            continue
        errors = []
        compare_range(actual_wb[sheet], golden_wb[sheet], item["range"], tolerance, errors)
        if errors:
            result["accuracy"].extend("ERROR:" + msg for msg in errors[:40])
            if len(errors) > 40:
                result["accuracy"].append(f"ERROR:还有 {len(errors) - 40} 个单元格不匹配")
        else:
            result["accuracy"].append(f"{sheet}!{item['range']} 与标准结果一致")

        formula_total = 0
        formula_bad = 0
        for row in golden_wb[sheet][item["range"]]:
            for golden_cell in row:
                if isinstance(golden_cell.value, str) and golden_cell.value.startswith("="):
                    formula_total += 1
                    actual_value = actual_wb[sheet][golden_cell.coordinate].value
                    if normalize_formula(actual_value) != normalize_formula(golden_cell.value):
                        formula_bad += 1
        if formula_total:
            if formula_bad:
                result["formula"].append(
                    f"ERROR:{sheet}!{item['range']} 有 {formula_bad}/{formula_total} 个公式不正确"
                )
            else:
                result["formula"].append(f"{sheet} 的 {formula_total} 个关键公式正确")

    for sheet, rows in config.get("header_rows", {}).items():
        if sheet not in actual_wb.sheetnames:
            continue
        actual_ws, golden_ws = actual_wb[sheet], golden_wb[sheet]
        for row_idx in rows:
            for col_idx in range(1, golden_ws.max_column + 1):
                golden_cell = golden_ws.cell(row_idx, col_idx)
                if golden_cell.value is None:
                    continue
                actual_cell = actual_ws.cell(row_idx, col_idx)
                if actual_cell.font.bold != golden_cell.font.bold:
                    result["format"].append(
                        f"ERROR:{sheet}!{actual_cell.coordinate} 表头粗体样式不正确"
                    )
                if color_rgb(actual_cell.fill.fgColor) != color_rgb(golden_cell.fill.fgColor):
                    result["format"].append(
                        f"ERROR:{sheet}!{actual_cell.coordinate} 表头填充色不正确"
                    )
        if not any(x.startswith("ERROR:") and x.startswith(f"ERROR:{sheet}!") for x in result["format"]):
            result["format"].append(f"{sheet} 表头样式正确")

    for sheet, expected in config.get("freeze_panes", {}).items():
        if sheet in actual_wb.sheetnames:
            value = actual_wb[sheet].freeze_panes
            value = value.coordinate if hasattr(value, "coordinate") else value
            if value != expected:
                result["format"].append(f"ERROR:{sheet} 冻结窗格应为 {expected}，实际为 {value}")

    for sheet, expected_count in config.get("min_charts", {}).items():
        if sheet in actual_wb.sheetnames:
            actual_charts = actual_wb[sheet]._charts
            golden_charts = golden_wb[sheet]._charts
            if len(actual_charts) < expected_count:
                result["format"].append(
                    f"ERROR:{sheet} 至少需要 {expected_count} 张可编辑图表，实际 {len(actual_charts)}"
                )
            elif [chart_type(c) for c in actual_charts[:expected_count]] != [
                chart_type(c) for c in golden_charts[:expected_count]
            ]:
                result["format"].append(f"ERROR:{sheet} 图表类型与要求不一致")
            else:
                result["format"].append(f"{sheet} 包含 {len(actual_charts)} 张类型正确的可编辑图表")

    for sheet, state in config.get("sheet_states", {}).items():
        if sheet in actual_wb.sheetnames and actual_wb[sheet].sheet_state != state:
            result["format"].append(f"ERROR:{sheet} 的工作表状态应为 {state}")

    errors = scan_errors(actual_wb)
    if errors:
        result["quality"].extend("ERROR:检测到表格错误 " + item for item in errors[:20])
    else:
        result["quality"].append("未检测到 Excel 错误标记")

    nonempty = sum(
        1 for ws in actual_wb.worksheets for row in ws.iter_rows() for cell in row if cell.value not in (None, "")
    )
    if nonempty < config.get("min_nonempty_cells", 10):
        result["quality"].append(f"ERROR:工作簿内容过少，仅 {nonempty} 个非空单元格")
    return result


def check_docx(candidate, golden, config):
    result = {"integrity": [], "structure": [], "accuracy": [], "formula": [], "format": [], "quality": []}
    try:
        actual = Document(candidate)
        expected = Document(golden)
        result["integrity"].append("Word 报告可正常打开")
    except Exception as exc:
        result["integrity"].append(f"ERROR:Word 报告无法打开: {exc}")
        return result

    actual_text = "\n".join(p.text for p in actual.paragraphs)
    expected_text = "\n".join(p.text for p in expected.paragraphs)
    for token in config.get("required_text", []):
        if token not in actual_text:
            result["structure"].append(f"ERROR:Word 报告缺少文本: {token}")
        else:
            result["structure"].append(f"Word 报告包含: {token}")

    for token in config.get("golden_numeric_tokens", []):
        if token not in actual_text or token not in expected_text:
            result["accuracy"].append(f"ERROR:Word 报告缺少关键数字: {token}")
        else:
            result["accuracy"].append(f"Word 报告关键数字正确: {token}")

    heading_count = sum(1 for p in actual.paragraphs if p.style and p.style.name.startswith("Heading"))
    if heading_count < config.get("min_headings", 3):
        result["format"].append(f"ERROR:Word 报告标题层级不完整，仅 {heading_count} 个标题")
    else:
        result["format"].append(f"Word 报告包含 {heading_count} 个标题")

    if len(actual.paragraphs) < config.get("min_paragraphs", 10):
        result["quality"].append("ERROR:Word 报告段落过少")
    else:
        result["quality"].append(f"Word 报告包含 {len(actual.paragraphs)} 个段落")
    if len(actual.tables) < config.get("min_tables", 1):
        result["structure"].append("ERROR:Word 报告缺少 KPI 表格")
    return result


def has_error(messages):
    return any(item.startswith("ERROR:") or ": ERROR:" in item for item in messages)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--candidate-dir", required=True)
    parser.add_argument("--golden-dir", required=True)
    parser.add_argument("--checks", required=True)
    parser.add_argument("--result-json")
    args = parser.parse_args()

    candidate_dir = Path(args.candidate_dir)
    golden_dir = Path(args.golden_dir)
    checks = json.loads(Path(args.checks).read_text(encoding="utf-8"))
    grouped = {key: [] for key in ["integrity", "structure", "accuracy", "formula", "format", "quality"]}

    for file_config in checks["files"]:
        name = file_config["name"]
        candidate = candidate_dir / name
        golden = golden_dir / name
        if not candidate.is_file():
            grouped["integrity"].append(f"ERROR:缺少必需文件 {name}")
            continue
        if candidate.stat().st_size < 500:
            grouped["integrity"].append(f"ERROR:{name} 文件过小，疑似空壳")
            continue
        if file_config["type"] == "xlsx":
            partial = check_xlsx(candidate, golden, file_config, checks.get("numeric_tolerance", 0.01))
        else:
            partial = check_docx(candidate, golden, file_config)
        for key, items in partial.items():
            grouped[key].extend(f"{name}: {item}" for item in items)

    criteria = [
        {"id": "R001", "group": "integrity", "weight": 2},
        {"id": "R002", "group": "structure", "weight": 2},
        {"id": "R003", "group": "accuracy", "weight": 4},
        {"id": "R004", "group": "formula", "weight": 3},
        {"id": "R005", "group": "format", "weight": 2},
        {"id": "R006", "group": "quality", "weight": 1},
    ]
    scored = []
    earned = 0
    total = sum(item["weight"] for item in criteria)
    for item in criteria:
        messages = grouped[item["group"]]
        passed = bool(messages) and not has_error(messages)
        if item["group"] == "formula" and not messages:
            # 没有指定公式的任务不因“无公式”失分。
            passed = True
            messages = ["本任务无额外强制公式检查"]
        if passed:
            earned += item["weight"]
        scored.append(
            {
                "id": item["id"],
                "passed": passed,
                "weight": item["weight"],
                "evidence": messages[:12],
            }
        )

    penalty = has_error(grouped["quality"]) or has_error(grouped["integrity"])
    if penalty:
        earned = max(0, earned - 4)
    score = earned / total if total else 0
    hurdle_passed = scored[0]["passed"]
    passed = hurdle_passed and score >= checks.get("pass_threshold", 0.8)
    result = {
        "task_id": checks["task_id"],
        "passed": passed,
        "score": round(score, 6),
        "pass_threshold": checks.get("pass_threshold", 0.8),
        "criteria": scored,
        "penalty": {"id": "R007", "triggered": penalty, "weight": -4},
    }
    payload = json.dumps(result, ensure_ascii=False, indent=2)
    print(payload)
    if args.result_json:
        Path(args.result_json).write_text(payload + "\n", encoding="utf-8")
    raise SystemExit(0 if passed else 1)


if __name__ == "__main__":
    main()
'''


def create_simple_workbook(path: Path, sheet_name: str, headers: list[str], rows: list[list]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    for row in rows:
        ws.append(row)
    finish_sheet(ws)
    add_excel_table(ws, "InputTable")
    save_workbook(wb, path)


def input_document_to_table(root: Path) -> None:
    document = Document()
    title = document.add_heading("行政设备询价会议纪要", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("内部编号：SYN-RFQ-2026-08；币种：人民币；报价均为含税价。")
    document.add_paragraph("会议时间：2026-08-18；报价有效期：30 天。")
    document.add_heading("报价记录", level=1)
    records = [
        (1, "杭州青杉办公设备有限公司", "MON-27", "27 英寸显示器", 12, 1268.00, 7),
        (2, "杭州青杉办公设备有限公司", "DOCK-8", "8 合 1 扩展坞", 20, 438.00, 5),
        (3, "苏州北辰数码科技有限公司", "CHAIR-E2", "人体工学椅 E2", 30, 786.00, 12),
        (4, "苏州北辰数码科技有限公司", "MON-27", "27 英寸显示器", 12, 1235.00, 10),
        (5, "宁波海河办公集成有限公司", "DOCK-8", "8 合 1 扩展坞", 20, 452.00, 3),
        (6, "宁波海河办公集成有限公司", "CHAIR-E2", "人体工学椅 E2", 30, 768.00, 15),
    ]
    for record in records:
        document.add_paragraph("|".join(str(value) for value in record))
    document.add_paragraph("字段顺序：报价序号|供应商|物料编码|物料名称|数量|含税单价|交付天数。")
    document.add_heading("决策规则", level=1)
    document.add_paragraph("同一物料先比较含税总额；本轮不考虑跨物料打包折扣。")
    path = root / "data/input_files/供应商报价纪要.docx"
    document.save(path)


def input_data_cleaning(root: Path) -> None:
    csv_path = root / "data/input_files/客户线索_待清洗.csv"
    rows = [
        ["L001", " 张敏 ", "138-0013-8000", "ZHANG.MIN@EXAMPLE.COM ", "浙江", "2026/08/01", "¥120,000", "是", "2026-08-01 09:10"],
        ["L002", "李雷", "139 1000 2000", " li.lei@example.com", "北京市", "2026-08-02", "8万", "Y", "2026-08-02 10:00"],
        ["L003", "王芳", "13700001111", "wang.fang@example.com", "沪", "2026.08.02", "45000", "true", "2026-08-02 10:30"],
        ["L003", "王芳", "137-0000-2222", "wang.fang@example.com", "上海", "2026.08.02", "5.2万", "是", "2026-08-04 15:45"],
        ["L004", "陈杰", "1360000333", "chen.jie@example", "广东", "2026/08/03", "30000", "否", "2026-08-03 11:20"],
        ["L005", "赵静", "+86 135 0000 4444", "zhao.jing@example.com", "粤", "2026/08/04", " 98,500 ", "YES", "2026-08-04 13:00"],
        ["L006", "周宁", "13400005555", "zhou.ning@example.com", "江苏省", "2026-08-04", "N/A", "N", "2026-08-04 14:20"],
        ["L007", "孙浩", "13300006666", "sun.hao@example.com", "浙", "2026-08-05", "72000", "1", "2026-08-05 08:40"],
        ["L008", "郑英", "13200007777", "ZHENG.YING@EXAMPLE.COM", "沪", "2026/08/05", "15万", "0", "2026-08-05 09:25"],
        ["L009", "何明", "13100008888", "he.ming @example.com", "未知", "2026/08/06", "66000", "是", "2026-08-06 16:10"],
        ["L010", "吴桐", "13000009999", "wu.tong@example.com", "京", "2026-08-07", "110000", "yes", "2026-08-07 12:00"],
    ]
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["lead_id", "name", "phone", "email", "province", "created_at", "budget", "consent", "updated_at"])
        writer.writerows(rows)

    mapping = [
        ["北京", "北京市"], ["北京市", "北京市"], ["京", "北京市"],
        ["上海", "上海市"], ["上海市", "上海市"], ["沪", "上海市"],
        ["浙江", "浙江省"], ["浙江省", "浙江省"], ["浙", "浙江省"],
        ["广东", "广东省"], ["广东省", "广东省"], ["粤", "广东省"],
        ["江苏", "江苏省"], ["江苏省", "江苏省"], ["苏", "江苏省"],
    ]
    create_simple_workbook(root / "data/input_files/省份映射.xlsx", "映射", ["别名", "标准省份"], mapping)


def input_table_transformation(root: Path) -> None:
    rows = [
        ["华东", "A 型净化器", 126000, 84, 138000, 90, 151500, 97],
        ["华东", "B 型加湿器", 86000, 172, 92500, 181, 101000, 194],
        ["华南", "A 型净化器", 112500, 75, 121500, 81, 133500, 86],
        ["华南", "B 型加湿器", 79000, 158, 84500, 166, 93000, 180],
        ["华北", "A 型净化器", 103500, 69, 114000, 76, 120000, 78],
        ["华北", "B 型加湿器", 71000, 142, 75500, 149, 82500, 160],
    ]
    headers = ["区域", "产品", "2026-01销售额", "2026-01销量", "2026-02销售额", "2026-02销量", "2026-03销售额", "2026-03销量"]
    create_simple_workbook(root / "data/input_files/区域月度销售_宽表.xlsx", "月度销售", headers, rows)


def input_formula_authoring(root: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "订单明细"
    ws.append(["订单号", "订单日期", "产品", "数量", "含税单价", "折扣率", "单位成本", "销售额", "成本", "毛利", "毛利率"])
    rows = [
        ["O-260801", date(2026, 8, 1), "Pro 订阅", 5, 1200, 0.05, 760],
        ["O-260802", date(2026, 8, 2), "Basic 订阅", 12, 350, 0, 210],
        ["O-260803", date(2026, 8, 3), "Team 订阅", 8, 890, 0.10, 560],
        ["O-260804", date(2026, 8, 5), "Pro 订阅", 3, 1200, 0, 760],
        ["O-260805", date(2026, 8, 8), "Basic 订阅", 20, 350, 0.08, 210],
        ["O-260806", date(2026, 8, 13), "Team 订阅", 6, 890, 0.05, 560],
        ["O-260807", date(2026, 8, 18), "Pro 订阅", 9, 1200, 0.12, 760],
        ["O-260808", date(2026, 8, 21), "Basic 订阅", 15, 350, 0.03, 210],
    ]
    for base in rows:
        ws.append(base + [None, None, None, None])
    finish_sheet(ws)
    for row in range(2, ws.max_row + 1):
        ws[f"B{row}"].number_format = "yyyy-mm-dd"
        ws[f"F{row}"].number_format = "0.00%"
        for col in "EGHIJ":
            ws[f"{col}{row}"].number_format = '¥#,##0.00'
        ws[f"K{row}"].number_format = "0.00%"
    summary = wb.create_sheet("经营摘要")
    summary.append(["指标", "数值"])
    for metric in ["订单数", "总销售额", "总成本", "总毛利", "整体毛利率"]:
        summary.append([metric, None])
    finish_sheet(summary)
    save_workbook(wb, root / "data/input_files/订单明细_待计算.xlsx")


def input_model_building(root: Path) -> None:
    wb = Workbook()
    history = wb.active
    history.title = "历史数据"
    history.append(["指标", "2024实际", "2025实际"])
    for row in [
        ["期末客户数", 820, 1040], ["月 ARPU", 465, 492], ["年流失率", 0.13, 0.115],
        ["销售成本率", 0.24, 0.225], ["薪酬", 2380000, 2980000], ["营销费用", 1260000, 1580000],
        ["其他费用", 820000, 960000],
    ]:
        history.append(row)
    finish_sheet(history)
    assumptions = wb.create_sheet("经营假设")
    assumptions.append(["假设", 2026, 2027, 2028])
    for row in [
        ["新增客户", 360, 420, 500], ["年流失率", 0.105, 0.095, 0.09], ["月 ARPU", 510, 530, 555],
        ["销售成本率", 0.215, 0.205, 0.195], ["薪酬", 3450000, 3850000, 4320000],
        ["营销费用", 1820000, 2050000, 2350000], ["其他费用", 1080000, 1180000, 1300000],
    ]:
        assumptions.append(row)
    finish_sheet(assumptions)
    for ws in [history, assumptions]:
        for row in range(2, ws.max_row + 1):
            if "率" in str(ws[f"A{row}"].value):
                for col in range(2, ws.max_column + 1):
                    ws.cell(row, col).number_format = "0.00%"
    save_workbook(wb, root / "data/input_files/SaaS经营假设与历史.xlsx")


def input_scenario_analysis(root: Path) -> None:
    wb = Workbook()
    base = wb.active
    base.title = "基础参数"
    base.append(["参数", "数值", "单位"])
    for row in [
        ["基准年包裹量", 120000, "票"], ["基准单票价格", 12.5, "元"],
        ["单票变动成本", 7.2, "元"], ["年固定成本", 380000, "元"], ["价格弹性", -0.8, "倍"]
    ]:
        base.append(row)
    finish_sheet(base)
    scenario = wb.create_sheet("情景参数")
    scenario.append(["情景", "调价幅度"])
    for row in [["悲观", -0.08], ["基准", 0], ["乐观", 0.10]]:
        scenario.append(row)
    finish_sheet(scenario)
    for row in range(2, 5):
        scenario[f"B{row}"].number_format = "0.00%"
    save_workbook(wb, root / "data/input_files/物流定价假设.xlsx")


def retail_transactions() -> list[list]:
    channels = ["门店", "电商", "经销商"]
    categories = ["净化器", "加湿器", "耗材"]
    regions = ["华东", "华南"]
    rows = []
    seq = 1
    for month in range(1, 5):
        for channel_index, channel in enumerate(channels):
            for category_index, category in enumerate(categories):
                region = regions[(channel_index + category_index + month) % 2]
                quantity = 12 + month * 3 + channel_index * 2 + category_index * 4
                unit_price = [1650, 620, 180][category_index]
                discount = 1 - 0.02 * channel_index
                revenue = round(quantity * unit_price * discount, 2)
                cost = round(revenue * [0.61, 0.57, 0.48][category_index], 2)
                rows.append([f"TX-{seq:04d}", date(2026, month, 8 + channel_index), region, channel, category, quantity, revenue, cost])
                seq += 1
    return rows


def input_pivot_analysis(root: Path) -> None:
    headers = ["交易号", "日期", "区域", "渠道", "品类", "数量", "销售额", "成本"]
    create_simple_workbook(root / "data/input_files/零售交易明细.xlsx", "交易", headers, retail_transactions())


def input_charting(root: Path) -> None:
    rows = [
        ["W01", date(2026, 6, 1), 1240, 310, 86], ["W02", date(2026, 6, 8), 1310, 345, 94],
        ["W03", date(2026, 6, 15), 1285, 338, 101], ["W04", date(2026, 6, 22), 1420, 372, 110],
        ["W05", date(2026, 6, 29), 1510, 410, 126], ["W06", date(2026, 7, 6), 1495, 422, 131],
        ["W07", date(2026, 7, 13), 1610, 455, 144], ["W08", date(2026, 7, 20), 1705, 492, 160],
    ]
    headers = ["周次", "周开始日", "线索数", "商机数", "订单数"]
    create_simple_workbook(root / "data/input_files/增长漏斗周报.xlsx", "原始数据", headers, rows)


def sales_orders() -> list[list]:
    regions = ["华东", "华南", "华北"]
    products = [("A 型净化器", 1680, 980), ("B 型加湿器", 650, 360), ("替换滤芯", 210, 95)]
    rows = []
    for index in range(1, 61):
        region = regions[(index - 1) % 3]
        product, price, cost = products[(index + 1) % 3]
        month = (index - 1) // 15 + 1
        quantity = 4 + (index * 3) % 17
        rows.append([f"SO-{index:04d}", date(2026, month, 1 + (index % 24)), region, product, quantity, price, cost])
    return rows


def input_dashboard(root: Path) -> None:
    wb = Workbook()
    orders = wb.active
    orders.title = "订单"
    orders.append(["订单号", "日期", "区域", "产品", "数量", "单价", "单位成本"])
    for row in sales_orders():
        orders.append(row)
    finish_sheet(orders)
    target = wb.create_sheet("销售目标")
    target.append(["区域", "2026Q1 目标", "2026Q2 目标"])
    for row in [["华东", 265000, 285000], ["华南", 248000, 268000], ["华北", 232000, 252000]]:
        target.append(row)
    finish_sheet(target)
    save_workbook(wb, root / "data/input_files/销售订单与目标.xlsx")


def input_template_completion(root: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "岗位需求"
    ws.merge_cells("A1:G1")
    ws["A1"] = "2026 年第四季度人员编制与招聘计划"
    ws["A1"].font = Font(name="微软雅黑", size=16, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32
    ws["A2"] = "更新日期：2026-08-24；编制口径：全职员工。"
    ws.merge_cells("A2:G2")
    headers = ["部门", "岗位", "目标编制", "现有人数", "招聘缺口", "期望到岗日", "优先级"]
    for col, value in enumerate(headers, 1):
        ws.cell(3, col, value)
    style_header(ws, 3)
    for row in range(4, 12):
        for col in range(1, 8):
            ws.cell(row, col).border = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = "A3:G11"
    widths = [16, 20, 12, 12, 12, 14, 12]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    note = wb.create_sheet("说明")
    note.append(["规则", "内容"])
    note.append(["缺口", "目标编制-现有人数，不小于 0"])
    note.append(["高优先级", "缺口>2 或期望到岗日早于 2026-10-01"])
    finish_sheet(note)
    note.sheet_state = "hidden"
    save_workbook(wb, root / "data/input_files/人员编制计划模板.xlsx")
    with (root / "data/input_files/招聘需求.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["部门", "岗位", "目标编制", "现有人数", "期望到岗日"])
        writer.writerows([
            ["研发部", "后端工程师", 12, 8, "2026-09-20"], ["研发部", "测试工程师", 6, 5, "2026-10-15"],
            ["产品部", "产品经理", 7, 5, "2026-09-25"], ["销售部", "大客户经理", 15, 11, "2026-10-20"],
            ["销售部", "售前顾问", 8, 7, "2026-11-01"], ["客服部", "客户成功经理", 10, 7, "2026-09-18"],
            ["财务部", "FP&A 分析师", 4, 3, "2026-10-10"], ["人力资源部", "招聘专员", 5, 4, "2026-09-30"],
        ])


def input_spreadsheet_repair(root: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "部门预算"
    ws.append(["部门", "Q1实际", "Q2实际", "Q3实际", "Q4实际", "年度合计", "获批预算", "差异"])
    data = [
        ["研发部", 820000, 910000, 960000, 1030000, "=SUM(B2:E2)", 3900000, "=G2-F2"],
        ["销售部", 610000, 680000, 720000, 795000, "=#REF!", 2900000, "=F3-G3"],
        ["客服部", 330000, 350000, 375000, 392000, "=SUM(B4:D4)", 1500000, "=G4-F4"],
        ["人力资源部", 280000, 315000, 305000, 338000, "=SUM(B5:E5)", 1300000, "=#REF!"],
        ["行政部", 210000, 225000, 248000, 260000, "=SUM(B6:E6)", 980000, "=G6-F6"],
        ["合计", "=SUM(B2:B6)", "=SUM(C2:C5)", "=SUM(D2:D6)", "=SUM(E2:E6)", "=SUM(F2:F6)", "=SUM(G2:G6)", "=G7-F7"],
    ]
    for row in data:
        ws.append(row)
    finish_sheet(ws)
    # 故意损坏一个表头样式。
    ws["D1"].fill = PatternFill("solid", fgColor="FFFFFF")
    ws["D1"].font = Font(name="微软雅黑", size=10, bold=False, color="000000")
    summary = wb.create_sheet("汇总")
    summary.append(["部门", "年度实际", "获批预算", "差异"])
    for idx, dept in enumerate(["研发部", "销售部", "客服部", "人力资源部", "行政部"], 2):
        summary.append([dept, f"='部门预算'!F{idx}", f"='部门预算'!G{idx}", "=#REF!"])
    finish_sheet(summary)
    save_workbook(wb, root / "data/input_files/年度预算_公式损坏.xlsx")


def input_spreadsheet_audit(root: Path) -> None:
    headers = ["交易编号", "日期", "发票号", "供应商", "合同号", "费用类别", "金额", "审批人"]
    rows = [
        ["T001", date(2026, 7, 1), "INV-7001", "恒通设备", "C-EQP", "设备", 78000, "林峰"],
        ["T002", date(2026, 7, 2), "INV-7002", "恒通设备", "C-EQP", "设备", 92000, "林峰"],
        ["T003", date(2026, 7, 3), "INV-7003", "建诚工程", "C-CON", "施工", 135000, "王宁"],
        ["T004", date(2026, 7, 4), "INV-7004", "建诚工程", "C-CON", "施工", 168000, "王宁"],
        ["T005", date(2026, 7, 5), "INV-7005", "优联服务", "C-SVC", "服务", 38000, "陈娜"],
        ["T006", date(2026, 7, 6), "INV-7006", "优联服务", "C-SVC", "服务", 42000, ""],
        ["T007", date(2026, 7, 8), "INV-7007", "恒通设备", "C-EQP", "设备", 126000, "林峰"],
        ["T008", date(2026, 7, 9), "INV-7008", "建诚工程", "C-CON", "施工", 99000, "王宁"],
        ["T009", date(2026, 7, 10), "INV-7009", "优联服务", "C-SVC", "服务", -6500, "陈娜"],
        ["T010", date(2026, 7, 11), "INV-7002", "恒通设备", "C-EQP", "设备", 92000, "林峰"],
        ["T011", date(2026, 7, 12), "INV-7011", "建诚工程", "C-CON", "施工", 188000, "王宁"],
        ["T012", date(2026, 7, 13), "INV-7012", "优联服务", "C-SVC", "服务", 28000, "陈娜"],
    ]
    create_simple_workbook(root / "data/input_files/项目成本台账_待审计.xlsx", "成本台账", headers, rows)
    limits = [
        ["C-EQP", "设备", 100000, 360000],
        ["C-CON", "施工", 180000, 550000],
        ["C-SVC", "服务", 50000, 120000],
    ]
    create_simple_workbook(root / "data/input_files/采购合同限额.xlsx", "合同限额", ["合同号", "费用类别", "单笔限额", "合同总额"], limits)


def input_multi_workbook(root: Path) -> None:
    headers = ["订单号", "订单日期", "SKU", "数量", "成交单价", "更新时间"]
    region_rows = {
        "华东": [
            ["B001", date(2026, 8, 1), "SKU-A", 12, 1580, "2026-08-02 09:00"],
            ["B002", date(2026, 8, 2), "SKU-B", 25, 620, "2026-08-02 14:00"],
            ["B003", date(2026, 8, 3), "SKU-C", 40, 198, "2026-08-03 15:00"],
            ["B004", date(2026, 8, 4), "SKU-A", 8, 1600, "2026-08-04 17:00"],
        ],
        "华南": [
            ["B005", date(2026, 8, 2), "SKU-B", 18, 610, "2026-08-02 12:00"],
            ["B006", date(2026, 8, 3), "SKU-A", 10, 1560, "2026-08-03 16:00"],
            ["B003", date(2026, 8, 3), "SKU-C", 42, 195, "2026-08-05 10:00"],
            ["B007", date(2026, 8, 5), "SKU-X", 6, 880, "2026-08-05 18:00"],
        ],
        "华北": [
            ["B008", date(2026, 8, 1), "SKU-C", 35, 205, "2026-08-01 13:00"],
            ["B009", date(2026, 8, 4), "SKU-B", 22, 625, "2026-08-04 11:00"],
            ["B004", date(2026, 8, 4), "SKU-A", 9, 1590, "2026-08-06 09:30"],
            ["B010", date(2026, 8, 6), "SKU-C", 50, 190, "2026-08-06 12:00"],
        ],
    }
    for region, rows in region_rows.items():
        create_simple_workbook(root / f"data/input_files/{region}订单.xlsx", "订单", headers, rows)
    master = [["SKU-A", "净化器", 980], ["SKU-B", "加湿器", 360], ["SKU-C", "耗材", 95]]
    create_simple_workbook(root / "data/input_files/产品主数据.xlsx", "产品", ["SKU", "品类", "标准成本"], master)


def service_months() -> list[list]:
    values = [
        (1, 8200, 0.942, 4.58, 18, 1280000, 760000),
        (2, 7900, 0.951, 4.61, 16, 1265000, 742000),
        (3, 8600, 0.936, 4.52, 22, 1320000, 785000),
        (4, 9100, 0.928, 4.46, 27, 1350000, 820000),
        (5, 9400, 0.917, 4.39, 31, 1375000, 848000),
        (6, 9700, 0.905, 4.32, 36, 1380000, 875000),
        (7, 10100, 0.912, 4.35, 33, 1410000, 892000),
        (8, 9800, 0.926, 4.43, 28, 1425000, 880000),
        (9, 9300, 0.938, 4.49, 24, 1440000, 865000),
        (10, 9000, 0.947, 4.55, 20, 1460000, 850000),
        (11, 8800, 0.955, 4.62, 17, 1490000, 842000),
        (12, 10500, 0.931, 4.50, 26, 1560000, 930000),
    ]
    return [[f"2026-{month:02d}", tickets, sla, csat, churn, revenue, cost] for month, tickets, sla, csat, churn, revenue, cost in values]


def input_analysis_to_report(root: Path) -> None:
    headers = ["月份", "工单量", "SLA达成率", "CSAT", "流失客户数", "收入", "服务成本"]
    create_simple_workbook(root / "data/input_files/客户服务月度数据.xlsx", "月度数据", headers, service_months())


INPUT_BUILDERS = {
    "document_to_table": input_document_to_table,
    "data_cleaning": input_data_cleaning,
    "table_transformation": input_table_transformation,
    "formula_authoring": input_formula_authoring,
    "model_building": input_model_building,
    "scenario_analysis": input_scenario_analysis,
    "pivot_analysis": input_pivot_analysis,
    "charting": input_charting,
    "dashboard": input_dashboard,
    "template_completion": input_template_completion,
    "spreadsheet_repair": input_spreadsheet_repair,
    "spreadsheet_audit": input_spreadsheet_audit,
    "multi_workbook_integration": input_multi_workbook,
    "analysis_to_report": input_analysis_to_report,
}


SOLVER_BODIES = {
    "document_to_table": r'''
def solve(input_dir, output_dir):
    document = Document(input_dir / "供应商报价纪要.docx")
    records = []
    for paragraph in document.paragraphs:
        parts = [part.strip() for part in paragraph.text.split("|")]
        if len(parts) == 7 and parts[0].isdigit():
            records.append({
                "seq": int(parts[0]), "vendor": parts[1], "code": parts[2], "name": parts[3],
                "quantity": int(parts[4]), "unit_price": float(parts[5]), "days": int(parts[6]),
            })
    if len(records) != 6:
        raise ValueError(f"期望 6 条报价，实际读取 {len(records)} 条")
    for item in records:
        item["total"] = round(item["quantity"] * item["unit_price"], 2)
    records.sort(key=lambda item: (item["code"], item["total"], item["seq"]))

    wb = Workbook()
    detail = wb.active
    detail.title = "报价明细"
    detail.append(["报价序号", "供应商", "物料编码", "物料名称", "数量", "含税单价", "含税总额", "交付天数"])
    for item in records:
        detail.append([item["seq"], item["vendor"], item["code"], item["name"], item["quantity"], item["unit_price"], item["total"], item["days"]])
    finish(detail)
    money(detail, ["F", "G"])
    table(detail, "QuoteDetails")

    vendor_summary = defaultdict(lambda: [0, 0.0, []])
    for item in records:
        summary = vendor_summary[item["vendor"]]
        summary[0] += 1
        summary[1] += item["total"]
        summary[2].append(item["days"])
    summary_ws = wb.create_sheet("供应商汇总")
    summary_ws.append(["供应商", "报价条数", "报价金额合计", "平均交付天数"])
    for vendor in sorted(vendor_summary):
        count, total, days = vendor_summary[vendor]
        summary_ws.append([vendor, count, round(total, 2), round(sum(days) / len(days), 2)])
    finish(summary_ws)
    money(summary_ws, ["C"])
    table(summary_ws, "VendorSummary")

    best_by_code = {}
    for item in records:
        best_by_code.setdefault(item["code"], item)
    recommendation = wb.create_sheet("最低价推荐")
    recommendation.append(["物料编码", "物料名称", "推荐供应商", "数量", "含税单价", "含税总额", "交付天数"])
    for code in sorted(best_by_code):
        item = best_by_code[code]
        recommendation.append([code, item["name"], item["vendor"], item["quantity"], item["unit_price"], item["total"], item["days"]])
    recommendation.append(["推荐合计", None, None, None, None, sum(item["total"] for item in best_by_code.values()), None])
    finish(recommendation)
    money(recommendation, ["E", "F"])
    recommendation[f"A{recommendation.max_row}"].font = Font(bold=True, color=WHITE)
    recommendation[f"A{recommendation.max_row}"].fill = PatternFill("solid", fgColor=NAVY)
    recommendation[f"F{recommendation.max_row}"].font = Font(bold=True)
    save(wb, output_dir / "采购报价结构化.xlsx")
''',
    "data_cleaning": r'''
def parse_budget(raw):
    text = str(raw).strip().replace("¥", "").replace(",", "")
    if not text or text.upper() in {"N/A", "NA", "NONE"}:
        return None
    multiplier = 10000 if text.endswith("万") else 1
    text = text[:-1] if text.endswith("万") else text
    try:
        return round(float(text) * multiplier, 2)
    except ValueError:
        return None


def solve(input_dir, output_dir):
    mapping_wb = load_workbook(input_dir / "省份映射.xlsx", data_only=True)
    mapping = {str(row[0]).strip(): str(row[1]).strip() for row in mapping_wb.active.iter_rows(min_row=2, values_only=True)}
    records = {}
    input_count = 0
    with (input_dir / "客户线索_待清洗.csv").open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            input_count += 1
            key = row["lead_id"].strip()
            if key not in records or row["updated_at"] > records[key]["updated_at"]:
                records[key] = row

    cleaned = []
    for key in sorted(records):
        row = records[key]
        name = row["name"].strip()
        phone = re.sub(r"\D", "", row["phone"])
        if len(phone) == 13 and phone.startswith("86"):
            phone = phone[2:]
        email = re.sub(r"\s+", "", row["email"]).lower()
        province = mapping.get(row["province"].strip(), "未知")
        budget = parse_budget(row["budget"])
        consent = "是" if row["consent"].strip().lower() in {"是", "y", "yes", "true", "1"} else "否"
        problems = []
        if not re.fullmatch(r"1[3-9]\d{9}", phone):
            problems.append("手机号无效")
        if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
            problems.append("邮箱无效")
        if province == "未知":
            problems.append("省份未匹配")
        if budget is None or budget < 0:
            problems.append("预算无效")
        cleaned.append([
            key, name, phone, email, province, row["created_at"].replace("/", "-").replace(".", "-"),
            budget, consent, row["updated_at"], "通过" if not problems else "；".join(problems),
        ])

    wb = Workbook()
    ws = wb.active
    ws.title = "清洗结果"
    ws.append(["lead_id", "姓名", "手机号", "邮箱", "标准省份", "创建日期", "预算（元）", "同意营销", "最后更新时间", "校验状态"])
    for row in cleaned:
        ws.append(row)
    finish(ws)
    money(ws, ["G"])
    table(ws, "CleanLeads")
    ws.conditional_formatting.add(
        f"J2:J{ws.max_row}", CellIsRule(operator="notEqual", formula=['"通过"'], fill=PatternFill("solid", fgColor="F4CCCC"))
    )

    summary = wb.create_sheet("质量汇总")
    pass_count = sum(1 for row in cleaned if row[-1] == "通过")
    summary.append(["指标", "数值"])
    for row in [
        ["输入记录数", input_count], ["去重后记录数", len(cleaned)], ["重复记录数", input_count - len(cleaned)],
        ["校验通过数", pass_count], ["校验不通过数", len(cleaned) - pass_count], ["通过率", pass_count / len(cleaned)],
    ]:
        summary.append(row)
    finish(summary)
    summary["B7"].number_format = "0.00%"

    province_ws = wb.create_sheet("省份统计")
    counts = defaultdict(int)
    for row in cleaned:
        counts[row[4]] += 1
    province_ws.append(["标准省份", "线索数"])
    for province in sorted(counts):
        province_ws.append([province, counts[province]])
    finish(province_ws)
    save(wb, output_dir / "客户线索清洗结果.xlsx")
''',
    "table_transformation": r'''
def solve(input_dir, output_dir):
    source = load_workbook(input_dir / "区域月度销售_宽表.xlsx", data_only=True).active
    months = ["2026-01", "2026-02", "2026-03"]
    long_rows = []
    for row in source.iter_rows(min_row=2, values_only=True):
        region, product = row[:2]
        for index, month in enumerate(months):
            sales, quantity = row[2 + index * 2], row[3 + index * 2]
            long_rows.append([region, product, month, sales, quantity, round(sales / quantity, 2)])
    long_rows.sort(key=lambda row: (row[0], row[1], row[2]))

    wb = Workbook()
    detail = wb.active
    detail.title = "销售长表"
    detail.append(["区域", "产品", "月份", "销售额", "销量", "客单价"])
    for row in long_rows:
        detail.append(row)
    finish(detail)
    money(detail, ["D", "F"])
    table(detail, "LongSales")

    def add_summary(name, key_index):
        ws = wb.create_sheet(name)
        grouped = defaultdict(lambda: [0.0, 0])
        for row in long_rows:
            grouped[row[key_index]][0] += row[3]
            grouped[row[key_index]][1] += row[4]
        ws.append(["分组", "销售额", "销量", "加权客单价"])
        for key in sorted(grouped):
            sales, quantity = grouped[key]
            ws.append([key, sales, quantity, round(sales / quantity, 2)])
        finish(ws)
        money(ws, ["B", "D"])
        return ws

    add_summary("区域汇总", 0)
    add_summary("产品汇总", 1)
    save(wb, output_dir / "区域月度销售长表.xlsx")
''',
    "formula_authoring": r'''
def solve(input_dir, output_dir):
    wb = load_workbook(input_dir / "订单明细_待计算.xlsx", data_only=False)
    ws = wb["订单明细"]
    for row in range(2, ws.max_row + 1):
        ws[f"H{row}"] = f"=D{row}*E{row}*(1-F{row})"
        ws[f"I{row}"] = f"=D{row}*G{row}"
        ws[f"J{row}"] = f"=H{row}-I{row}"
        ws[f"K{row}"] = f'=IFERROR(J{row}/H{row},0)'
        for col in "EGHIJ":
            ws[f"{col}{row}"].number_format = '¥#,##0.00'
        ws[f"K{row}"].number_format = "0.00%"
    finish(ws)
    summary = wb["经营摘要"]
    last = ws.max_row
    formulas = [
        f"=COUNTA('订单明细'!A2:A{last})", f"=SUM('订单明细'!H2:H{last})",
        f"=SUM('订单明细'!I2:I{last})", f"=SUM('订单明细'!J2:J{last})", "=IFERROR(B5/B2,0)",
    ]
    for row, formula in enumerate(formulas, 2):
        summary[f"B{row}"] = formula
    for row in range(3, 6):
        summary[f"B{row}"].number_format = '¥#,##0.00'
    summary["B6"].number_format = "0.00%"
    finish(summary)
    save(wb, output_dir / "订单毛利计算.xlsx")
''',
    "model_building": r'''
def solve(input_dir, output_dir):
    wb = load_workbook(input_dir / "SaaS经营假设与历史.xlsx", data_only=False)
    model = wb.create_sheet("三年模型")
    model.merge_cells("A1:D1")
    model["A1"] = "SaaS 2026–2028 三年经营模型"
    model["A1"].font = Font(name="微软雅黑", size=15, bold=True, color=WHITE)
    model["A1"].fill = PatternFill("solid", fgColor=NAVY)
    model["A1"].alignment = Alignment(horizontal="center")
    model.append(["指标", 2026, 2027, 2028])
    metrics = ["期初客户数", "新增客户", "流失客户", "期末客户数", "平均客户数", "月 ARPU", "订阅收入", "销售成本率", "销售成本", "毛利", "毛利率", "薪酬", "营销费用", "其他费用", "EBITDA", "EBITDA 率"]
    for metric in metrics:
        model.append([metric, None, None, None])
    for col, letter in enumerate(["B", "C", "D"], 2):
        assumption_col = get_column_letter(col)
        model[f"{letter}3"] = "='历史数据'!C2" if letter == "B" else f"={get_column_letter(col-1)}6"
        model[f"{letter}4"] = f"='经营假设'!{assumption_col}2"
        model[f"{letter}5"] = f"={letter}3*'经营假设'!{assumption_col}3"
        model[f"{letter}6"] = f"={letter}3+{letter}4-{letter}5"
        model[f"{letter}7"] = f"=AVERAGE({letter}3,{letter}6)"
        model[f"{letter}8"] = f"='经营假设'!{assumption_col}4"
        model[f"{letter}9"] = f"={letter}7*{letter}8*12"
        model[f"{letter}10"] = f"='经营假设'!{assumption_col}5"
        model[f"{letter}11"] = f"={letter}9*{letter}10"
        model[f"{letter}12"] = f"={letter}9-{letter}11"
        model[f"{letter}13"] = f"=IFERROR({letter}12/{letter}9,0)"
        model[f"{letter}14"] = f"='经营假设'!{assumption_col}6"
        model[f"{letter}15"] = f"='经营假设'!{assumption_col}7"
        model[f"{letter}16"] = f"='经营假设'!{assumption_col}8"
        model[f"{letter}17"] = f"={letter}12-SUM({letter}14:{letter}16)"
        model[f"{letter}18"] = f"=IFERROR({letter}17/{letter}9,0)"
    finish(model, header_row=2, freeze="B3", autofilter=False)
    for row in [8, 9, 11, 12, 14, 15, 16, 17]:
        money(model, ["B", "C", "D"], row, row)
    for row in [10, 13, 18]:
        percent(model, ["B", "C", "D"], row, row)
    for row in [4, 5, 8, 10, 14, 15, 16]:
        for col in range(2, 5):
            model.cell(row, col).fill = PatternFill("solid", fgColor=LIGHT_BLUE)

    summary = wb.create_sheet("指标摘要")
    summary.append(["指标", "数值"])
    summary.append(["收入 CAGR（2026–2028）", "=IFERROR(('三年模型'!D9/'三年模型'!B9)^(1/2)-1,0)"])
    summary.append(["2028 年毛利率", "='三年模型'!D13"])
    summary.append(["2028 年 EBITDA 率", "='三年模型'!D18"])
    finish(summary)
    for row in range(2, 5):
        summary[f"B{row}"].number_format = "0.00%"
    save(wb, output_dir / "SaaS三年经营模型.xlsx")
''',
    "scenario_analysis": r'''
def solve(input_dir, output_dir):
    wb = load_workbook(input_dir / "物流定价假设.xlsx", data_only=False)
    scenario = wb.create_sheet("情景分析")
    scenario.append(["情景", "调价幅度", "调整后价格", "预计包裹量", "收入", "变动成本", "贡献毛利", "营业利润"])
    for row in range(2, 5):
        scenario[f"A{row}"] = f"='情景参数'!A{row}"
        scenario[f"B{row}"] = f"='情景参数'!B{row}"
        scenario[f"C{row}"] = f"='基础参数'!B3*(1+B{row})"
        scenario[f"D{row}"] = f"='基础参数'!B2*(1+'基础参数'!B6*B{row})"
        scenario[f"E{row}"] = f"=C{row}*D{row}"
        scenario[f"F{row}"] = f"=D{row}*'基础参数'!B4"
        scenario[f"G{row}"] = f"=E{row}-F{row}"
        scenario[f"H{row}"] = f"=G{row}-'基础参数'!B5"
    finish(scenario)
    percent(scenario, ["B"], 2, 4)
    money(scenario, ["C", "E", "F", "G", "H"], 2, 4)

    matrix = wb.create_sheet("敏感性矩阵")
    matrix["A1"] = "营业利润敏感性：调价幅度 × 额外业务量变动"
    matrix.merge_cells("A1:F1")
    matrix.append(["额外业务量变动", -0.10, -0.05, 0, 0.05, 0.10])
    for row, volume_change in enumerate([-0.10, 0, 0.10], 3):
        matrix[f"A{row}"] = volume_change
        for col in range(2, 7):
            letter = get_column_letter(col)
            matrix[f"{letter}{row}"] = (
                f"=('基础参数'!B2*(1+'基础参数'!B6*{letter}$2)*(1+$A{row}))*"
                f"('基础参数'!B3*(1+{letter}$2)-'基础参数'!B4)-'基础参数'!B5"
            )
            matrix[f"{letter}{row}"].number_format = '¥#,##0.00;[Red]-¥#,##0.00'
    header(matrix, 2)
    body(matrix, 3)
    for cell in matrix[2][1:]:
        cell.number_format = "0%"
    for row in range(3, 6):
        matrix[f"A{row}"].number_format = "0%"
    matrix.freeze_panes = "B3"
    matrix.sheet_view.showGridLines = False
    matrix.conditional_formatting.add("B3:F5", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"))
    for col in range(1, 7):
        matrix.column_dimensions[get_column_letter(col)].width = 18
    save(wb, output_dir / "物流定价情景分析.xlsx")
''',
    "pivot_analysis": r'''
def solve(input_dir, output_dir):
    source_wb = load_workbook(input_dir / "零售交易明细.xlsx", data_only=True)
    source = source_wb.active
    rows = list(source.iter_rows(min_row=2, values_only=True))
    wb = Workbook()
    detail = wb.active
    detail.title = "交易明细"
    detail.append([cell.value for cell in source[1]] + ["毛利"])
    for row in rows:
        detail.append(list(row) + [row[6] - row[7]])
    finish(detail)
    money(detail, ["G", "H", "I"])
    table(detail, "RetailTransactions")

    channels = sorted({row[3] for row in rows})
    categories = sorted({row[4] for row in rows})
    revenue = defaultdict(float)
    gross = defaultdict(float)
    for row in rows:
        revenue[(row[3], row[4])] += row[6]
        gross[row[3]] += row[6] - row[7]
    pivot = wb.create_sheet("透视分析")
    pivot.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(categories) + 3)
    pivot["A1"] = "渠道 × 品类销售额透视式汇总"
    pivot["A1"].font = Font(size=14, bold=True, color=WHITE)
    pivot["A1"].fill = PatternFill("solid", fgColor=NAVY)
    pivot["A1"].alignment = Alignment(horizontal="center")
    pivot.append(["渠道"] + categories + ["行合计", "毛利率"])
    for channel in channels:
        values = [revenue[(channel, category)] for category in categories]
        total = sum(values)
        pivot.append([channel] + values + [total, gross[channel] / total])
    column_totals = [sum(revenue[(channel, category)] for channel in channels) for category in categories]
    grand_total = sum(column_totals)
    pivot.append(["列合计"] + column_totals + [grand_total, sum(gross.values()) / grand_total])
    header(pivot, 2)
    body(pivot, 3)
    pivot.freeze_panes = "B3"
    for col in range(2, len(categories) + 3):
        money(pivot, [get_column_letter(col)], 3, pivot.max_row)
    percent(pivot, [get_column_letter(len(categories) + 3)], 3, pivot.max_row)
    pivot.sheet_view.showGridLines = False
    for col in range(1, pivot.max_column + 1):
        pivot.column_dimensions[get_column_letter(col)].width = 16

    trend = wb.create_sheet("月度趋势")
    monthly = defaultdict(lambda: [0.0, 0.0])
    for row in rows:
        month = row[1].strftime("%Y-%m")
        monthly[month][0] += row[6]
        monthly[month][1] += row[6] - row[7]
    trend.append(["月份", "销售额", "毛利", "毛利率"])
    for month in sorted(monthly):
        sales, margin = monthly[month]
        trend.append([month, sales, margin, margin / sales])
    finish(trend)
    money(trend, ["B", "C"])
    percent(trend, ["D"])
    chart = BarChart()
    chart.type = "col"
    chart.title = "月度销售额趋势"
    chart.y_axis.title = "销售额（元）"
    chart.x_axis.title = "月份"
    chart.add_data(Reference(trend, min_col=2, min_row=1, max_row=trend.max_row), titles_from_data=True)
    chart.set_categories(Reference(trend, min_col=1, min_row=2, max_row=trend.max_row))
    chart.height = 8
    chart.width = 15
    trend.add_chart(chart, "F2")
    save(wb, output_dir / "零售渠道透视分析.xlsx")
''',
    "charting": r'''
def solve(input_dir, output_dir):
    source = load_workbook(input_dir / "增长漏斗周报.xlsx", data_only=True).active
    wb = Workbook()
    data = wb.active
    data.title = "周度数据"
    data.append(["周次", "周开始日", "线索数", "商机数", "订单数", "线索→商机转化率", "商机→订单转化率"])
    for row in source.iter_rows(min_row=2, values_only=True):
        data.append(list(row) + [row[3] / row[2], row[4] / row[3]])
    finish(data)
    percent(data, ["F", "G"])
    table(data, "GrowthFunnel")

    board = wb.create_sheet("图表看板")
    board.append(["指标", "数值"])
    total_leads = sum(data[f"C{r}"].value for r in range(2, data.max_row + 1))
    total_opps = sum(data[f"D{r}"].value for r in range(2, data.max_row + 1))
    total_orders = sum(data[f"E{r}"].value for r in range(2, data.max_row + 1))
    for row in [["累计线索", total_leads], ["累计商机", total_opps], ["累计订单", total_orders], ["整体线索转订单率", total_orders / total_leads]]:
        board.append(row)
    finish(board)
    board["B5"].number_format = "0.00%"

    line = LineChart()
    line.title = "周度漏斗规模趋势"
    line.y_axis.title = "数量"
    line.x_axis.title = "周次"
    line.add_data(Reference(data, min_col=3, max_col=5, min_row=1, max_row=data.max_row), titles_from_data=True)
    line.set_categories(Reference(data, min_col=1, min_row=2, max_row=data.max_row))
    line.height, line.width = 8, 15
    board.add_chart(line, "D2")
    bars = BarChart()
    bars.type = "col"
    bars.title = "周度转化率"
    bars.y_axis.title = "转化率"
    bars.y_axis.numFmt = "0%"
    bars.x_axis.title = "周次"
    bars.add_data(Reference(data, min_col=6, max_col=7, min_row=1, max_row=data.max_row), titles_from_data=True)
    bars.set_categories(Reference(data, min_col=1, min_row=2, max_row=data.max_row))
    bars.height, bars.width = 8, 15
    board.add_chart(bars, "D18")
    save(wb, output_dir / "增长漏斗图表.xlsx")
''',
    "dashboard": r'''
def solve(input_dir, output_dir):
    source_wb = load_workbook(input_dir / "销售订单与目标.xlsx", data_only=True)
    orders = list(source_wb["订单"].iter_rows(min_row=2, values_only=True))
    targets = {row[0]: row[1] + row[2] for row in source_wb["销售目标"].iter_rows(min_row=2, values_only=True)}
    wb = Workbook()
    detail = wb.active
    detail.title = "明细"
    detail.append(["订单号", "日期", "区域", "产品", "数量", "单价", "单位成本", "销售额", "成本", "毛利"])
    enriched = []
    for row in orders:
        sales = row[4] * row[5]
        cost = row[4] * row[6]
        item = list(row) + [sales, cost, sales - cost]
        enriched.append(item)
        detail.append(item)
    finish(detail)
    money(detail, ["F", "G", "H", "I", "J"])
    table(detail, "SalesDetails")

    monthly = defaultdict(lambda: [0.0, 0.0])
    regional = defaultdict(lambda: [0, 0.0, 0.0])
    product = defaultdict(lambda: [0.0, 0.0])
    for row in enriched:
        month = row[1].strftime("%Y-%m")
        monthly[month][0] += row[7]
        monthly[month][1] += row[9]
        regional[row[2]][0] += 1
        regional[row[2]][1] += row[7]
        regional[row[2]][2] += row[9]
        product[row[3]][0] += row[7]
        product[row[3]][1] += row[9]
    analysis = wb.create_sheet("分析")
    analysis.append(["月份", "销售额", "毛利", "毛利率", None, "区域", "订单数", "销售额", "毛利", "目标", "达成率", None, "产品", "销售额", "毛利", "毛利贡献率"])
    max_rows = max(len(monthly), len(regional), len(product))
    months, regions, products = sorted(monthly), sorted(regional), sorted(product)
    total_gross = sum(values[1] for values in product.values())
    for idx in range(max_rows):
        row = [None] * 16
        if idx < len(months):
            month = months[idx]
            row[0:4] = [month, monthly[month][0], monthly[month][1], monthly[month][1] / monthly[month][0]]
        if idx < len(regions):
            region = regions[idx]
            row[5:11] = [region, regional[region][0], regional[region][1], regional[region][2], targets[region], regional[region][1] / targets[region]]
        if idx < len(products):
            name = products[idx]
            row[12:16] = [name, product[name][0], product[name][1], product[name][1] / total_gross]
        analysis.append(row)
    header(analysis, 1)
    body(analysis, 2)
    analysis.freeze_panes = "A2"
    analysis.sheet_view.showGridLines = False
    for col in ["B", "C", "H", "I", "J", "N", "O"]:
        money(analysis, [col], 2, analysis.max_row)
    for col in ["D", "K", "P"]:
        percent(analysis, [col], 2, analysis.max_row)
    for col in range(1, 17):
        analysis.column_dimensions[get_column_letter(col)].width = 14

    dashboard = wb.create_sheet("仪表盘")
    dashboard.append(["2026 年销售运营仪表盘", None, None, None, None])
    dashboard.merge_cells("A1:E1")
    dashboard["A1"].font = Font(size=16, bold=True, color=WHITE)
    dashboard["A1"].fill = PatternFill("solid", fgColor=NAVY)
    dashboard["A1"].alignment = Alignment(horizontal="center")
    total_sales = sum(row[7] for row in enriched)
    total_cost = sum(row[8] for row in enriched)
    kpis = [
        ("A3", "总销售额", total_sales), ("B3", "总毛利", total_sales - total_cost),
        ("C3", "整体毛利率", (total_sales - total_cost) / total_sales), ("D3", "订单数", len(enriched)),
        ("E3", "目标达成率", total_sales / sum(targets.values())),
    ]
    for cell_ref, label, value in kpis:
        cell = dashboard[cell_ref]
        cell.value = f"{label}\n{value:,.2f}" if "率" not in label else f"{label}\n{value:.2%}"
        cell.font = Font(size=12, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    dashboard.row_dimensions[3].height = 50
    for col in range(1, 6):
        dashboard.column_dimensions[get_column_letter(col)].width = 22
    dashboard.sheet_view.showGridLines = False

    line = LineChart()
    line.title = "月度销售趋势"
    line.add_data(Reference(analysis, min_col=2, min_row=1, max_row=len(months) + 1), titles_from_data=True)
    line.set_categories(Reference(analysis, min_col=1, min_row=2, max_row=len(months) + 1))
    line.height, line.width = 7, 13
    dashboard.add_chart(line, "A6")
    region_chart = BarChart()
    region_chart.title = "区域目标达成率"
    region_chart.add_data(Reference(analysis, min_col=11, min_row=1, max_row=len(regions) + 1), titles_from_data=True)
    region_chart.set_categories(Reference(analysis, min_col=6, min_row=2, max_row=len(regions) + 1))
    region_chart.y_axis.numFmt = "0%"
    region_chart.height, region_chart.width = 7, 13
    dashboard.add_chart(region_chart, "H6")
    pie = PieChart()
    pie.title = "产品毛利贡献"
    pie.add_data(Reference(analysis, min_col=15, min_row=1, max_row=len(products) + 1), titles_from_data=True)
    pie.set_categories(Reference(analysis, min_col=13, min_row=2, max_row=len(products) + 1))
    pie.height, pie.width = 7, 13
    dashboard.add_chart(pie, "A21")
    dashboard.sheet_properties.pageSetUpPr.fitToPage = True
    dashboard.page_setup.orientation = "landscape"
    dashboard.page_setup.paperSize = dashboard.PAPERSIZE_A4
    dashboard.page_setup.fitToWidth = 1
    dashboard.page_setup.fitToHeight = 1
    dashboard.print_area = "A1:T38"
    dashboard.page_margins.left = 0.25
    dashboard.page_margins.right = 0.25
    dashboard.page_margins.top = 0.35
    dashboard.page_margins.bottom = 0.35
    save(wb, output_dir / "销售运营仪表盘.xlsx")
''',
    "template_completion": r'''
def solve(input_dir, output_dir):
    wb = load_workbook(input_dir / "人员编制计划模板.xlsx", data_only=False)
    ws = wb["岗位需求"]
    rows = []
    with (input_dir / "招聘需求.csv").open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    for index, item in enumerate(rows, 4):
        ws[f"A{index}"] = item["部门"]
        ws[f"B{index}"] = item["岗位"]
        ws[f"C{index}"] = int(item["目标编制"])
        ws[f"D{index}"] = int(item["现有人数"])
        ws[f"E{index}"] = f"=MAX(C{index}-D{index},0)"
        ws[f"F{index}"] = datetime.strptime(item["期望到岗日"], "%Y-%m-%d").date()
        ws[f"G{index}"] = f'=IF(OR(E{index}>2,F{index}<DATE(2026,10,1)),"高","常规")'
        ws[f"F{index}"].number_format = "yyyy-mm-dd"
    body(ws, 4)
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:G{3 + len(rows)}"
    ws.conditional_formatting.add(
        f"G4:G{3 + len(rows)}", CellIsRule(operator="equal", formula=['"高"'], fill=PatternFill("solid", fgColor="F4CCCC"))
    )
    summary = wb.create_sheet("部门汇总")
    summary.append(["部门", "目标编制", "现有人数", "招聘缺口"])
    departments = sorted({row["部门"] for row in rows})
    last = 3 + len(rows)
    for index, department in enumerate(departments, 2):
        summary[f"A{index}"] = department
        summary[f"B{index}"] = f'=SUMIF(\'岗位需求\'!A4:A{last},A{index},\'岗位需求\'!C4:C{last})'
        summary[f"C{index}"] = f'=SUMIF(\'岗位需求\'!A4:A{last},A{index},\'岗位需求\'!D4:D{last})'
        summary[f"D{index}"] = f'=SUMIF(\'岗位需求\'!A4:A{last},A{index},\'岗位需求\'!E4:E{last})'
    finish(summary)
    save(wb, output_dir / "人员编制计划完成版.xlsx")
''',
    "spreadsheet_repair": r'''
def solve(input_dir, output_dir):
    wb = load_workbook(input_dir / "年度预算_公式损坏.xlsx", data_only=False)
    budget = wb["部门预算"]
    summary = wb["汇总"]
    changes = []

    def repair(ws, cell_ref, new_value, reason):
        old = ws[cell_ref].value
        if old != new_value:
            changes.append([ws.title, cell_ref, old, new_value, reason])
            ws[cell_ref] = new_value

    for row in range(2, 7):
        repair(budget, f"F{row}", f"=SUM(B{row}:E{row})", "修复年度合计范围")
        repair(budget, f"H{row}", f"=G{row}-F{row}", "统一差异口径为获批预算-实际")
    for col in "BCDEFG":
        repair(budget, f"{col}7", f"=SUM({col}2:{col}6)", "修复合计行范围")
    repair(budget, "H7", "=G7-F7", "修复合计差异")
    for row in range(2, 7):
        repair(summary, f"B{row}", f"='部门预算'!F{row}", "修复年度实际引用")
        repair(summary, f"C{row}", f"='部门预算'!G{row}", "修复获批预算引用")
        repair(summary, f"D{row}", f"='部门预算'!H{row}", "修复差异引用")
    changes.append(["部门预算", "D1", "表头白色/非粗体", "深蓝底白色粗体", "恢复表头统一样式"])
    header(budget, 1)
    finish(budget)
    finish(summary)
    for ws in [budget, summary]:
        money(ws, [get_column_letter(col) for col in range(2, ws.max_column + 1)])

    log = wb.create_sheet("修复日志")
    log.append(["工作表", "单元格", "原公式/值", "修复后公式/值", "修复原因"])
    for row in changes:
        escaped = list(row)
        for index in [2, 3]:
            if isinstance(escaped[index], str) and escaped[index].startswith("="):
                escaped[index] = "'" + escaped[index]
        log.append(escaped)
    finish(log)
    save(wb, output_dir / "年度预算修复版.xlsx")
''',
    "spreadsheet_audit": r'''
def solve(input_dir, output_dir):
    ledger_wb = load_workbook(input_dir / "项目成本台账_待审计.xlsx", data_only=True)
    limit_wb = load_workbook(input_dir / "采购合同限额.xlsx", data_only=True)
    headers = [cell.value for cell in ledger_wb.active[1]]
    ledger = [dict(zip(headers, row)) for row in ledger_wb.active.iter_rows(min_row=2, values_only=True)]
    limits = {row[0]: {"category": row[1], "single": row[2], "total": row[3]} for row in limit_wb.active.iter_rows(min_row=2, values_only=True)}
    findings = []
    number = 1

    def add(risk, level, tx, evidence, suggestion, amount):
        nonlocal number
        findings.append([f"F-{number:03d}", risk, level, tx, evidence, suggestion, amount])
        number += 1

    by_invoice = defaultdict(list)
    by_contract = defaultdict(list)
    for item in ledger:
        by_invoice[item["发票号"]].append(item)
        by_contract[item["合同号"]].append(item)
        limit = limits[item["合同号"]]
        if item["金额"] > limit["single"]:
            add("单笔超限", "高", item["交易编号"], f"金额 {item['金额']:.2f} > 单笔限额 {limit['single']:.2f}", "追加授权审批或冲销超限部分", item["金额"])
        if not str(item["审批人"] or "").strip():
            add("缺失审批人", "中", item["交易编号"], "审批人字段为空", "补齐可核验的审批记录", abs(item["金额"]))
        if item["金额"] < 0:
            add("负数金额", "中", item["交易编号"], f"金额为 {item['金额']:.2f}", "核实是否为合法退款并关联原发票", abs(item["金额"]))
    for invoice, items in sorted(by_invoice.items()):
        if len(items) > 1:
            add("重复发票", "高", ",".join(item["交易编号"] for item in items), f"发票号 {invoice} 出现 {len(items)} 次", "暂停支付并核对原始票据", sum(abs(item["金额"]) for item in items))
    for contract, items in sorted(by_contract.items()):
        total = sum(item["金额"] for item in items)
        limit = limits[contract]["total"]
        if total > limit:
            add("合同累计超限", "高", ",".join(item["交易编号"] for item in items), f"合同 {contract} 累计 {total:.2f} > 限额 {limit:.2f}", "停止新增支付并发起合同变更审批", total - limit)

    wb = Workbook()
    original = wb.active
    original.title = "原始台账"
    original.append(headers)
    for item in ledger:
        original.append([item[h] for h in headers])
    finish(original)
    money(original, ["G"])
    table(original, "OriginalLedger")

    finding_ws = wb.create_sheet("审计发现")
    finding_ws.append(["发现编号", "风险类型", "风险级别", "交易编号", "证据", "建议", "涉及金额"])
    for row in findings:
        finding_ws.append(row)
    finish(finding_ws)
    money(finding_ws, ["G"])
    finding_ws.conditional_formatting.add(f"C2:C{finding_ws.max_row}", CellIsRule(operator="equal", formula=['"高"'], fill=PatternFill("solid", fgColor="F4CCCC")))
    table(finding_ws, "AuditFindings")

    summary = wb.create_sheet("审计汇总")
    grouped = defaultdict(lambda: [0, 0.0])
    for finding in findings:
        grouped[finding[1]][0] += 1
        grouped[finding[1]][1] += finding[6]
    summary.append(["风险类型", "发现数", "涉及金额"])
    for risk in sorted(grouped):
        summary.append([risk, grouped[risk][0], grouped[risk][1]])
    summary.append(["合计", len(findings), sum(item[6] for item in findings)])
    finish(summary)
    money(summary, ["C"])
    save(wb, output_dir / "项目成本审计.xlsx")
''',
    "multi_workbook_integration": r'''
def solve(input_dir, output_dir):
    master_wb = load_workbook(input_dir / "产品主数据.xlsx", data_only=True)
    master = {row[0]: (row[1], row[2]) for row in master_wb.active.iter_rows(min_row=2, values_only=True)}
    selected = {}
    discarded = []
    for region in ["华东", "华南", "华北"]:
        wb = load_workbook(input_dir / f"{region}订单.xlsx", data_only=True)
        for row in wb.active.iter_rows(min_row=2, values_only=True):
            record = {"order": row[0], "date": row[1], "sku": row[2], "quantity": row[3], "price": row[4], "updated": row[5], "region": region}
            previous = selected.get(record["order"])
            if previous is None or record["updated"] > previous["updated"]:
                if previous is not None:
                    discarded.append([previous["order"], previous["region"], previous["updated"], region, record["updated"], "保留更新时间最晚的记录"])
                selected[record["order"]] = record
            else:
                discarded.append([record["order"], region, record["updated"], previous["region"], previous["updated"], "保留更新时间最晚的记录"])

    wb = Workbook()
    detail = wb.active
    detail.title = "合并明细"
    detail.append(["订单号", "订单日期", "来源分公司", "SKU", "品类", "数量", "成交单价", "销售额", "标准单位成本", "标准成本", "标准毛利", "更新时间", "主数据状态"])
    enriched = []
    for order in sorted(selected):
        item = selected[order]
        category, unit_cost = master.get(item["sku"], ("未匹配", None))
        sales = item["quantity"] * item["price"]
        cost = item["quantity"] * unit_cost if unit_cost is not None else None
        gross = sales - cost if cost is not None else None
        row = [order, item["date"], item["region"], item["sku"], category, item["quantity"], item["price"], sales, unit_cost, cost, gross, item["updated"], "已匹配" if unit_cost is not None else "未匹配"]
        enriched.append(row)
        detail.append(row)
    finish(detail)
    money(detail, ["G", "H", "I", "J", "K"])
    detail.conditional_formatting.add(f"M2:M{detail.max_row}", CellIsRule(operator="equal", formula=['"未匹配"'], fill=PatternFill("solid", fgColor="F4CCCC")))
    table(detail, "IntegratedOrders")

    summary = wb.create_sheet("分公司汇总")
    grouped = defaultdict(lambda: [0, 0.0, 0.0])
    for row in enriched:
        grouped[row[2]][0] += 1
        grouped[row[2]][1] += row[7]
        if row[10] is not None:
            grouped[row[2]][2] += row[10]
    summary.append(["来源分公司", "订单数", "销售额", "标准毛利（已匹配 SKU）"])
    for region in sorted(grouped):
        summary.append([region] + grouped[region])
    finish(summary)
    money(summary, ["C", "D"])

    reconciliation = wb.create_sheet("对账清单")
    reconciliation.append(["订单号", "被舍弃来源", "被舍弃更新时间", "保留来源", "保留更新时间", "原因"])
    for row in sorted(discarded):
        reconciliation.append(row)
    finish(reconciliation)
    save(wb, output_dir / "分公司订单汇总.xlsx")
''',
    "analysis_to_report": r'''
def solve(input_dir, output_dir):
    source = load_workbook(input_dir / "客户服务月度数据.xlsx", data_only=True).active
    rows = list(source.iter_rows(min_row=2, values_only=True))
    wb = Workbook()
    data = wb.active
    data.title = "月度数据"
    data.append(["月份", "工单量", "SLA达成率", "CSAT", "流失客户数", "收入", "服务成本", "利润", "利润率"])
    for row in rows:
        profit = row[5] - row[6]
        data.append(list(row) + [profit, profit / row[5]])
    finish(data)
    percent(data, ["C", "I"])
    money(data, ["F", "G", "H"])
    table(data, "ServiceMonthly")

    total_tickets = sum(row[1] for row in rows)
    avg_sla = sum(row[2] for row in rows) / len(rows)
    weighted_csat = sum(row[3] * row[1] for row in rows) / total_tickets
    churn = sum(row[4] for row in rows)
    revenue = sum(row[5] for row in rows)
    cost = sum(row[6] for row in rows)
    profit = revenue - cost
    kpi = wb.create_sheet("KPI摘要")
    kpi.append(["指标", "数值"])
    for row in [
        ["全年工单量", total_tickets], ["平均 SLA 达成率", avg_sla], ["加权平均 CSAT", weighted_csat],
        ["流失客户数", churn], ["全年收入", revenue], ["全年服务成本", cost], ["全年利润", profit], ["全年利润率", profit / revenue],
    ]:
        kpi.append(row)
    finish(kpi)
    for cell in ["B3", "B9"]:
        kpi[cell].number_format = "0.00%"
    kpi["B4"].number_format = "0.00"
    money(kpi, ["B"], 6, 8)

    trend = wb.create_sheet("趋势分析")
    trend.append(["月份", "SLA达成率", "CSAT", "利润", "利润率"])
    for row in rows:
        monthly_profit = row[5] - row[6]
        trend.append([row[0], row[2], row[3], monthly_profit, monthly_profit / row[5]])
    finish(trend)
    percent(trend, ["B", "E"])
    money(trend, ["D"])
    quality_chart = LineChart()
    quality_chart.title = "SLA 与 CSAT 月度趋势"
    quality_chart.add_data(Reference(trend, min_col=2, max_col=3, min_row=1, max_row=trend.max_row), titles_from_data=True)
    quality_chart.set_categories(Reference(trend, min_col=1, min_row=2, max_row=trend.max_row))
    quality_chart.height, quality_chart.width = 8, 15
    trend.add_chart(quality_chart, "G2")
    profit_chart = BarChart()
    profit_chart.title = "月度利润趋势"
    profit_chart.add_data(Reference(trend, min_col=4, min_row=1, max_row=trend.max_row), titles_from_data=True)
    profit_chart.set_categories(Reference(trend, min_col=1, min_row=2, max_row=trend.max_row))
    profit_chart.height, profit_chart.width = 8, 15
    trend.add_chart(profit_chart, "G18")
    save(wb, output_dir / "客户服务经营分析.xlsx")

    lowest_sla = min(rows, key=lambda row: row[2])
    lowest_profit = min(rows, key=lambda row: row[5] - row[6])
    document = Document()
    title = document.add_heading("2026 年客户服务经营分析报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("数据口径：2026 年1–12月合成月度经营数据；本报告不包含外部市场推断。")
    document.add_heading("执行摘要", level=1)
    document.add_paragraph(
        f"全年处理 {total_tickets:,} 个工单，平均 SLA 达成率 {avg_sla:.2%}，"
        f"加权平均 CSAT {weighted_csat:.2f}。全年收入 ¥{revenue:,.2f}，服务成本 ¥{cost:,.2f}，"
        f"利润 ¥{profit:,.2f}，利润率 {profit/revenue:.2%}。"
    )
    table_doc = document.add_table(rows=1, cols=2)
    table_doc.style = "Table Grid"
    table_doc.rows[0].cells[0].text = "指标"
    table_doc.rows[0].cells[1].text = "数值"
    for label, value in [
        ("全年工单量", f"{total_tickets:,}"), ("平均 SLA", f"{avg_sla:.2%}"),
        ("加权 CSAT", f"{weighted_csat:.2f}"), ("流失客户数", f"{churn}"), ("全年利润", f"¥{profit:,.2f}"),
    ]:
        cells = table_doc.add_row().cells
        cells[0].text, cells[1].text = label, value
    document.add_heading("关键发现", level=1)
    document.add_paragraph(f"SLA 最低月份为 {lowest_sla[0]}，达成率 {lowest_sla[2]:.2%}；与全年平均值相比明显偏低。")
    document.add_paragraph(f"利润最低月份为 {lowest_profit[0]}，当月利润 ¥{lowest_profit[5]-lowest_profit[6]:,.2f}。")
    document.add_paragraph(f"全年流失客户 {churn} 个，SLA 较低的 5–7 月同时出现 CSAT 下探，应优先复盘人力排班和工单分流。")
    document.add_heading("改进建议", level=1)
    for text in [
        "1. 针对 5–7 月工单高峰，按周追踪未结工单和超时原因，设置 SLA 预警阈值。",
        "2. 将高频问题建立标准回复和分流规则，以缩短首次响应和解决时间。",
        "3. 对当月低分 CSAT 样本开展分层回访，每月向责任团队反馈前三类根因。",
        "4. 将月度利润与 SLA/CSAT 同时纳入经营会，防止单纯压降成本损害服务质量。",
    ]:
        document.add_paragraph(text)
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.font.name = "微软雅黑"
            run.font.size = Pt(10.5)
    document.save(output_dir / "客户服务经营分析报告.docx")
''',
}


SOLVER_CLI = r'''


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--output-dir", required=True)
    args = parser.parse_args()
    input_dir = Path(args.input_dir).resolve()
    output_dir = Path(args.output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    solve(input_dir, output_dir)


if __name__ == "__main__":
    main()
'''


def build_checks(spec: dict, root: Path) -> dict:
    files = []
    for output_name in spec["outputs"]:
        golden = root / "solution" / output_name
        suffix = golden.suffix.lower()
        if suffix == ".xlsx":
            wb = load_workbook(golden, data_only=False)
            ranges = []
            freeze_panes = {}
            min_charts = {}
            sheet_states = {}
            nonempty = 0
            for ws in wb.worksheets:
                if ws.max_row > 0 and ws.max_column > 0:
                    ranges.append({"sheet": ws.title, "range": ws.calculate_dimension()})
                freeze = ws.freeze_panes
                freeze = freeze.coordinate if hasattr(freeze, "coordinate") else freeze
                if freeze:
                    freeze_panes[ws.title] = freeze
                if ws._charts:
                    min_charts[ws.title] = len(ws._charts)
                sheet_states[ws.title] = ws.sheet_state
                nonempty += sum(
                    1 for row in ws.iter_rows() for cell in row if cell.value not in (None, "")
                )
            files.append(
                {
                    "name": output_name,
                    "type": "xlsx",
                    "sheets": wb.sheetnames,
                    "exact_sheet_order": True,
                    "ranges": ranges,
                    "header_rows": spec["header_rows"],
                    "freeze_panes": freeze_panes,
                    "min_charts": min_charts,
                    "sheet_states": sheet_states,
                    "min_nonempty_cells": max(10, int(nonempty * 0.8)),
                }
            )
        elif suffix == ".docx":
            workbook = load_workbook(root / "solution" / "客户服务经营分析.xlsx", data_only=False)
            kpi = workbook["KPI摘要"]
            numeric_tokens = [
                f"{int(kpi['B2'].value):,}",
                f"{kpi['B3'].value:.2%}",
                f"{kpi['B4'].value:.2f}",
                f"¥{kpi['B8'].value:,.2f}",
            ]
            files.append(
                {
                    "name": output_name,
                    "type": "docx",
                    "required_text": ["执行摘要", "关键发现", "改进建议", "SLA 最低月份", "利润最低月份"],
                    "golden_numeric_tokens": numeric_tokens,
                    "min_headings": 4,
                    "min_paragraphs": 10,
                    "min_tables": 1,
                }
            )
        else:
            raise ValueError(f"不支持的交付文件类型: {output_name}")
    return {
        "task_id": f"cowork-evo/{spec['task_id']}",
        "pass_threshold": 0.8,
        "numeric_tolerance": 0.01,
        "files": files,
    }


def build_provenance(spec: dict, root: Path) -> tuple[dict, dict]:
    inputs = []
    for path in sorted((root / "data" / "input_files").iterdir()):
        if not path.is_file():
            continue
        inputs.append(
            {
                "path": f"data/input_files/{path.name}",
                "sha256": sha256(path),
                "bytes": path.stat().st_size,
                "generated": True,
                "privacy": "原创合成数据，不含真实个人信息",
            }
        )
    manifest = {
        "task_id": f"cowork-evo/{spec['task_id']}",
        "generated_at": "2026-08-24T00:00:00+08:00",
        "generator": "tools/task_generation/generate_xlsx_tasks.py",
        "generation_method": "固定业务规则与确定性数值生成；未复制参考仓库数据",
        "inputs": inputs,
        "reference_repositories": REFERENCE_REPOSITORIES,
    }
    notes = {
        "task_family": spec["family"],
        "generation_method": manifest["generation_method"],
        "input_checksums": {item["path"]: item["sha256"] for item in inputs},
        "reference_repositories": REFERENCE_REPOSITORIES,
        "release_note": "仅借鉴公开仓库的题族、Rubric 和 Judge 设计；没有执行外部仓库脚本，也没有复制其数据包。",
    }
    provenance = {
        "provenance_id": f"{spec['task_id']}-synthetic-v1",
        "source_type": "synthetic",
        "source_uri": "https://github.com/RUCKBReasoning/SpreadsheetBench-2",
        "acquired_at": "2026-08-24T00:00:00+08:00",
        "license": "原创合成数据与任务文本；建议随 Benchmark 以 CC BY 4.0 发布，正式发布前复核",
        "privacy_status": "通过",
        "human_rewritten": True,
        "synthetic": True,
        "publishability": "public",
        "source_sha256": [item["sha256"] for item in inputs],
        "notes": json.dumps(notes, ensure_ascii=False, separators=(",", ":")),
    }
    return provenance, manifest


def build_task(spec: dict) -> None:
    root = make_base_task(spec)
    INPUT_BUILDERS[spec["family"]](root)

    solver_source = COMMON_SOLVER + "\n" + SOLVER_BODIES[spec["family"]] + SOLVER_CLI
    write_text(root / "solution" / "solve.py", solver_source, executable=True)
    write_text(root / "tests" / "judge.py", JUDGE_SOURCE, executable=True)
    write_text(root / "rubric.json", json.dumps(rubric_for(spec), ensure_ascii=False, indent=2) + "\n")

    completed = subprocess.run(
        [
            str(PYTHON),
            str(root / "solution" / "solve.py"),
            "--input-dir",
            str(root / "data" / "input_files"),
            "--output-dir",
            str(root / "solution"),
        ],
        cwd=REPO_ROOT,
        text=True,
        capture_output=True,
        check=False,
    )
    if completed.returncode != 0:
        raise RuntimeError(
            f"{spec['task_id']} Golden 生成失败\nSTDOUT:\n{completed.stdout}\nSTDERR:\n{completed.stderr}"
        )
    missing = [name for name in spec["outputs"] if not (root / "solution" / name).is_file()]
    if missing:
        raise RuntimeError(f"{spec['task_id']} 缺少 Golden: {missing}")

    checks = build_checks(spec, root)
    write_text(root / "tests" / "checks.json", json.dumps(checks, ensure_ascii=False, indent=2) + "\n")
    provenance, manifest = build_provenance(spec, root)
    write_text(root / "metadata" / "provenance.json", json.dumps(provenance, ensure_ascii=False, indent=2) + "\n")
    write_text(root / "metadata" / "input-manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--only", choices=[spec["family"] for spec in TASK_SPECS])
    args = parser.parse_args()
    selected = [spec for spec in TASK_SPECS if not args.only or spec["family"] == args.only]
    for spec in selected:
        build_task(spec)
    print(f"已生成 {len(selected)} 个 Excel 题包。")


if __name__ == "__main__":
    main()
